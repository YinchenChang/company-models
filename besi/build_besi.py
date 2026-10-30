# BESI（BE Semiconductor Industries，BESI.AS）收支與估值模型 v0.1.1（AI 半導體 Project 步驟 5：設備與材料）
# 用法：python3 besi/build_besi.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年（FY2026＝2026-01～12）。金額：百萬歐元（€M）；EPS：歐元。
# 樣板：hanmi/（鍵合設備；HBM 晶粒與混合鍵合比例口徑相同）、advantest/（設備股統一循環規則）
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
import cycle_position as CP

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
YRS = (2025, 2026, 2027, 2028, 2029)

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
SKH_F = only("skhynix/dist/*.xlsx"); DSC_F = only("disco/dist/*.xlsx"); EQ_F = only("markets/equipment/dist/*.xlsx")
TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)

ws_ = openpyxl.load_workbook(SKH_F, data_only=True)["Mkt"]
SHIP = {}   # HBM 出貨位元（EB），三情境各一列（markets/memory：需求 × SCEN_UNITS 情境顆數係數，與供給取小）
for r in ws_.iter_rows():
    lab = str(r[0].value or "")
    if lab.startswith("出貨＝MIN"):
        SHIP[lab[-2]] = ([r[2 + j].value for j in range(5)], f"{rel(SKH_F)}｜Mkt!C{r[0].row}:G{r[0].row}")
assert sorted(SHIP) == ["A", "B", "C"] and all(v is not None for s in SHIP for v in SHIP[s][0])

def in_keys(f, prefixes):
    out = {}
    for r in openpyxl.load_workbook(f, data_only=True)["In"].iter_rows(min_row=4):
        k = r[0].value
        if isinstance(k, str) and k.startswith(prefixes):
            out[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(f)}｜In!D{r[0].row}:F{r[0].row}")
    return out
DS = in_keys(DSC_F, ("den_", "stk_", "hbp_", "g_tail"))
# v0.7.1：晶粒容量、堆疊層數、混合鍵合占比改讀 Siliconomics M22（與 Disco 模型同一值）
for _k in [k for k in DS if k.startswith(("den_", "stk_", "hbp_"))]:
    DS[_k] = (SIL.vals("hbm_" + _k), f"siliconomics｜SIF_hbm_{_k}")
assert all(f"{p}_{y}" in DS for p in ("den", "stk", "hbp") for y in range(25, 31)) and "g_tail" in DS
TI = in_keys(TSMC_F, ("um28_", "um29_"))
assert all(f"um{y}_{s}" in TI for y in (28, 29) for s in SC)

ws_ = openpyxl.load_workbook(TSMC_F, data_only=True)["CoWoS"]
COW = None
for r in ws_.iter_rows():
    if r[0].value == "需求合計":
        COW = ([r[1 + j].value for j in range(4)], f"{rel(TSMC_F)}｜CoWoS!B{r[0].row}:E{r[0].row}")
assert COW and all(v is not None for v in COW[0])

ws_ = openpyxl.load_workbook(EQ_F, data_only=True)["Mkt"]
EQG = {}; sc_ = None   # 設備總額年增（markets/equipment Mkt 頁，情境 A／B／C）
for r in ws_.iter_rows(min_row=4, max_row=40):
    a = r[0].value
    if isinstance(a, str) and a.startswith("情境 ") and len(a.strip()) == 4:
        sc_ = a.strip()[-1]; continue
    if sc_ and isinstance(a, str) and a.strip() == "年增" and sc_ not in EQG:
        EQG[sc_] = ({y: r[2 + y - 2025].value for y in (2027, 2028, 2029)}, f"{rel(EQ_F)}｜Mkt!E{r[0].row}:G{r[0].row}")
assert sorted(EQG) == list(SC)
PEAK, PEAK_SRC = CP.peak_years(ROOT)   # 設備股共用循環位置規則：各情境見頂年度（markets/equipment）

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：百萬歐元（€M）；HBM 位元：EB；晶粒：百萬顆；CoWoS：千片。每個輸入附兩個來源（來源2 找不到者寫「未找到」）"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司（besi.com）只給下一季指引，不揭露混合鍵合營收、產品別營收與機台單價；實績取自公司新聞稿（GlobeNewswire 轉載）與財務資料頁"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)

sec("A. HBM 出貨位元（markets/memory，經 SK 海力士模型成品 Mkt 頁；情境 A／B／C 各一列，讀同字母；唯讀）")
for s in SC:
    for j, y in enumerate(YRS):
        inp(f"hs{s}_{y % 100}", f"HBM 出貨位元 {y}（EB；{s}）", "EB", None, SHIP[s][0][j], None, "Derived", SHIP[s][1],
            "markets/memory：出貨＝MIN(需求 × 情境顆數係數 SCEN_UNITS, 供給)；2028–29 隨情境，2025–27 三情境相同" if j == 0 else "同上")
v, src = DS["g_tail"]
inp("g_tail", "CY2030 需求相對 2029 的年增（市場層只到 2029；只用於 CY2029 的機台購買）", "%", v[0], v[1], v[2], "Assumed", src, "與 Disco、韓美模型的 g_tail 同值（讀 Disco 成品）")
sec("B. 晶粒、堆疊層數與混合鍵合比例（與 Disco、韓美模型同一組假設，讀 Disco 成品；唯讀）")
for y in range(25, 31):
    v, src = DS[f"den_{y}"]
    inp(f"den_{y}", f"DRAM 核心晶粒容量（組合平均）20{y}", "Gb／顆", v[0], v[1], v[2], "Assumed", src, "與韓美同一口徑：每顆核心晶粒鍵合一次；容量愈大，同樣位元的鍵合次數愈少" if y == 25 else "同上")
for y in range(25, 31):
    v, src = DS[f"stk_{y}"]
    inp(f"stk_{y}", f"HBM 平均堆疊層數 20{y}", "層", v[0], v[1], v[2], "Assumed", src, "層數愈高，每顆晶粒的鍵合精度要求與時間愈高（見 beta）" if y == 25 else "同上")
inp("beta", "堆疊層數對每顆晶粒鍵合時間的影響（每多 4 層，鍵合時間 +beta）", "x", 0.0, 0.10, 0.25, "Assumed", "與韓美模型同值（Assumed）：層數加權＝1＋beta×(層數 − 2025 層數)÷4", "第二來源：未找到")
for y in range(25, 31):
    v, src = DS[f"hbp_{y}"]
    inp(f"hbp_{y}", f"混合鍵合占 HBM 鍵合產能比例 20{y}（Disco 模型：低／基準／高）", "%", v[0], v[1], v[2], "Assumed", src,
        "與韓美模型同一比例（韓美用 1 − 此比例）。本模型：情境 A 取高值、B 取基準、C 取低值（對 BESI：A 最有利、C 最不利）；公司 Q3-26 法說：記憶體短缺使 HBM 混合鍵合導入延後、現行層數延長（Interested-party）" if y == 25 else "同上")
inp("dlt", "已裝機鍵合產能每年汰換／升級比例（新機需求＝產能增量＋汰換）", "%/年", 0.05, 0.10, 0.15, "Assumed", "與 Disco、韓美模型同值（Assumed）", "公司不揭露裝機量與機齡；精度由 100nm→50nm→25nm 換代（Q3-26 法說）")
sec("C. 邏輯混合鍵合（SoIC 等）：AI 先進封裝晶圓需求（台積電模型成品，基準路徑 × 同字母情境顆數係數；唯讀）")
for j, y in enumerate((2026, 2027, 2028, 2029)):
    inp(f"cow_{y % 100}", f"CoWoS 晶圓需求合計 {y}（基準路徑）", "千片", None, COW[0][j], None, "Derived", COW[1],
        "台積電模型：加速器顆數 ÷ 每片封裝數＋其他用途；本模型以此代表 AI 邏輯先進封裝的規模（Analogy：SoIC 晶圓與 CoWoS 同步）" if j == 0 else "同上")
for s in SC:
    for y in (28, 29):
        v, src = TI[f"um{y}_{s}"]
        inp(f"um{y}_{s}", f"情境 {s}：20{y} 顆數 ÷ 基準（台積電模型情境顆數係數）", "x", None, v[1], None, "Derived", src, "與 markets/memory SCEN_UNITS 同一套；2026–27 無情境差異" if (s, y) == ("A", 28) else "同上")
ATT = {"A": (0.08, 0.11, 0.18, 0.26, 0.34), "B": (0.08, 0.11, 0.15, 0.20, 0.25), "C": (0.08, 0.11, 0.12, 0.14, 0.16)}
for s in SC:
    for j, y in enumerate(range(2026, 2031)):
        b = ATT[s][j]
        inp(f"att{s}_{y % 100}", f"混合鍵合（SoIC 類）占 AI 先進封裝晶圓比例 {y}（{s}）", "%", round(b * (0.7 if j >= 2 else 1), 4), b, round(b * (1.3 if j >= 2 else 1), 4), "Assumed",
            "無直接資料：公司稱邏輯混合鍵合加速（AMD、資料中心與消費性 CPU 新品，Q2-26 新聞稿；美國邏輯客戶 2025 年裝機、明年擴產，Q3-26 法說；邏輯大客戶 6 條整合產線 30 台，Q4-25 新聞稿）" if (s, j) == ("A", 0) else "同上",
            "2026–27 三情境相同（2027 產能的機台多在 2026 下單，視為已知）；2028 起 A 導入最快、C 最慢；低／高＝2028 起 ×0.7／×1.3。比例的水準被校準係數 k 吸收，只有變化才影響營收" if j == 0 else "—")
sec("D. 混合鍵合機台（公司不揭露單價與產能；皆為 Assumed／Interested-party）")
inp("uph", "混合鍵合機有效產出（晶粒／小時）", "顆／時", 600, 1000, 1500, "Assumed", "公司產品頁：Datacon 8800 CHAMEO ultra plus AC 乾循環最高 2,000 CPH（Interested-party）；實際產出低於乾循環", "第二來源：未找到")
inp("util", "機台稼動率（全年小時 × 比例）", "%", 0.6, 0.7, 0.8, "Assumed", "無：後段設備常見 60–80%", "—")
inp("thr", "每台年產能（百萬顆晶粒）＝uph × 8760 × util ÷ 10⁶（Derived）", "百萬顆／台", None, "=uph*8760*util/1000000", None, "Derived", "—", "—")
inp("asp", "混合鍵合機單價（HBM 用；只用於 HBM 段）", "€M／台", 2.0, 2.5, 3.5, "Assumed", "公司不揭露（Q3-26 法說：未給 ASP）；隱含檢查：FY2026 混合鍵合營收 ÷ 單價 vs 累計訂單 150＋（2025 年底）→ 215＋台（Q3-26），見 Recon", "第二來源：未找到")
for s, b in (("A", 0.60), ("B", 0.50), ("C", 0.35)):
    inp(f"shm_{s}", f"BESI 占 HBM 混合鍵合機採購（{s}）", "%", b - 0.10, b, b + 0.10, "Assumed",
        "無份額資料：競爭者韓美（2027 年混合鍵合工廠）、韓華 Semitech、SEMES（三星）、ASMPT；BESI 與應用材料整合方案（Interested-party：公司稱 HBM4E／HBM5 採用混合鍵合）" if s == "A" else "同上",
        "A 最有利、C 最不利（Assumed）")
for s, b in (("A", 1.0), ("B", 0.95), ("C", 0.80)):
    inp(f"shl_{s}", f"BESI 邏輯混合鍵合份額，相對 FY2026（{s}）", "x", b - 0.10, b, min(b + 0.10, 1.1), "Assumed",
        "FY2026 份額已含在校準係數 k；之後 ASMPT、Kulicke & Soffa、韓美、ASML（公司稱非競爭者，Q3-26 法說，Interested-party）可能搶份額" if s == "A" else "同上", "A 守住、C 流失兩成（Assumed）")
sec("E. FY2026 錨點（前三季實績＋第四季公司指引）與混合鍵合營收拆分")
inp("q4g", "第四季營收相對第三季（公司指引 −5%～＋5%）", "%", -0.05, 0.0, 0.05, "Interested-party", "公司 Q3-26 新聞稿（2026-10-22）：第四季營收持平 ±5%（交貨時點，非需求轉弱）", "investing.com 簡報摘要、gurufocus 法說摘要同")
inp("q4gm", "第四季毛利率（公司指引 63–65%）", "%", 0.63, 0.64, 0.65, "Interested-party", "同上", "第三季 65.2%")
inp("q4opx", "第四季營業費用相對第三季（指引 +5%～+10%，含策略檢討顧問費）", "%", 0.05, 0.075, 0.10, "Interested-party", "同上：16 週策略檢討，2027-01 初結論", "—")
inp("q4fin", "第四季財務費用淨額", "€M", 4, 6, 8, "Assumed", "前三季 5.2、6.3、約 6.1（第三季由淨利 ÷（1 − 稅率）回推）", "2031 年到期 4.5% 優先票據 €350M")
inp("rev26", "FY2026 營收錨點（Derived）", "€M", None, "=q1_rev+q2_rev+q3_rev+q3_rev*(1+q4g)", None, "Derived", "前三季實績＋第三季 ×（1＋q4g）", "共識 €1,002M（marketscreener）")
inp("gp26", "FY2026 毛利錨點（Derived）", "€M", None, "=q1_gp+q2_gp+q3_gp+q3_rev*(1+q4g)*q4gm", None, "Derived", "—", "—")
inp("opx26", "FY2026 營業費用錨點（Derived）", "€M", None, "=q1_opx+q2_opx+q3_opx+q3_opx*(1+q4opx)", None, "Derived", "公司說法換算 FY2026 營業費用約 €223M（investing.com 簡報摘要）", "—")
inp("oi26", "FY2026 營業利益錨點（Derived）", "€M", None, "=gp26-opx26", None, "Derived", "—", "—")
inp("ni26", "FY2026 淨利錨點（Derived）", "€M", None, "=q1_ni+q2_ni+q3_ni+(q3_rev*(1+q4g)*q4gm-q3_opx*(1+q4opx)-q4fin)*(1-tax)", None, "Derived", "前三季實績＋第四季（營業利益 − 財務費用）×（1 − 稅率）", "—")
inp("eps26", "FY2026 EPS 錨點（Derived）", "€", None, "=q1_eps+q2_eps+q3_eps+(ni26-q1_ni-q2_ni-q3_ni)/shares", None, "Derived", "前三季 EPS 實績＋第四季淨利 ÷ 股數", "共識 4.33（marketscreener）／4.42（stockanalysis）")
inp("hb25", "FY2025 混合鍵合營收", "€M", 90, 120, 160, "Assumed", "公司不揭露。FT（經 tspasemiconductor 轉引）：2023 年 €36M；Morningstar（2025-02-21）：混合鍵合營收自 2023 年三倍（Interested-party）", "2025 年底累計訂單 150＋台（2024 年底未找到）；只用於 FY2025 拆分顯示")
inp("hb_sh26", "FY2026 混合鍵合營收占營收", "%", 0.18, 0.25, 0.35, "Assumed", "公司：2026 年前三季混合鍵合訂單超過 2025 全年兩倍（Q3-26 法說，Interested-party）；AI 相關占系統訂單約 60%（H1-26）", "FT（舊預測）：2026 年 €476M（約三分之一）；Sens 列低／高")
inp("hbm26", "FY2026 HBM 用混合鍵合營收（量產前）", "€M", None, 0, None, "Assumed", "公司：記憶體短缺使 HBM 混合鍵合導入延後（Q3-26 法說）；試產機台併入邏輯與其他（校準係數 k）", "—")
sec("F. 主流業務（傳統貼合機、封裝、電鍍、TC 鍵合；FY2027 起隨設備市場同字母情境）")
for s in SC:
    for y in (2027, 2028, 2029):
        inp(f"eqg{s}_{y % 100}", f"設備總額年增 {y}（{s}）", "%", None, EQG[s][0][y], None, "Derived", EQG[s][1], "markets/equipment Mkt 頁同字母情境（不自訂行業路徑）" if (s, y) == ("A", 2027) else "同上")
inp("el", "主流業務營收對設備市場年增的彈性", "x", 0.5, 1.0, 1.5, "Assumed", "無：公司稱主流市場（行動、車用、工業）2026 年復甦；車用約占 10%（Q3-26 法說）", "FY2025 主流業務 −7%（設備市場 +）時彈性不穩定，Sens 列 0.5／1.5")
sec("G. 毛利、營業費用、稅、股數")
inp("gm_prem", "混合鍵合相對主流業務的毛利率溢價（組合效果）", "pt", 0.0, 0.08, 0.15, "Assumed", "公司：第三、四季毛利率因產品組合變動（Q2-26、Q3-26 新聞稿）；長期毛利率指引 64–68%（投資人日 2026-06-19，Interested-party）", "FY2027 起毛利率＝FY2026 ＋ 溢價 ×（混合鍵合占比 − FY2026 占比）")
inp("gm_adj_C", "情境 C：FY2028 起毛利率額外調整（價格競爭）", "pt", -0.04, -0.02, 0.0, "Assumed", "無", "—")
inp("phi", "營業費用中變動部分占比", "%", 0.25, 0.40, 0.60, "Assumed", "FY2024 €200M、FY2025 €201M（營收 −3%）、FY2026 約 €223M（營收 +69%）：費用多為固定", "營業費用＝固定（1−phi）× FY2026 ＋ 變動 phi × FY2026 ÷ FY2026 營收 × 營收")
inp("opex_fix", "營業費用固定部分", "€M", None, "=(1-phi)*opx26", None, "Derived", "—", "—")
inp("opex_var", "營業費用變動部分占營收", "%", None, "=phi*opx26/rev26", None, "Derived", "—", "—")
inp("fin", "財務費用淨額（FY2027 起）", "€M", 10, 20, 30, "Assumed", "FY2025 €19.1M；2026 前三季約 €17.6M；€350M 4.5% 票據利息約 €16M，現金與存款 €592M 利息收入抵銷一部分", "—")
inp("tax", "有效稅率", "%", 0.12, 0.135, 0.16, "Verified", "公司：Q2-26 13.1%、Q3-26 13.3%（新聞稿；荷蘭創新盒）", "FY2025 14.5%（22.3 ÷ 153.9）")
inp("shares", "股數（Q3-26 淨利 ÷ EPS，加權平均）", "百萬股", None, "=q3_ni/q3_eps", None, "Derived", "107.2 ÷ 1.33 ≈ 80.6M；Q2-26 基本加權 80.04M、稀釋 80.93M（新聞稿）", "流通股 80.82M（stockanalysis）；可轉債 €175M 已轉換；每年 €100M 回購（2026-10-27 起）本版不建模")
sec("H. 估值、共識與對帳來源")
for s in SC:
    inp(f"pk_{s}", f"情境 {s}：設備市場見頂年度（日曆年；循環位置規則用）", "年", None, PEAK[s], None, "Derived", PEAK_SRC[s], "設備股共用規則（tools/cycle_position.py）：FY(t) 早於見頂年＝中段、等於＝見頂、晚於＝谷底；本公司會計年度＝日曆年" if s == "A" else "同上")
inp("px", "股價（stockanalysis 2026-10-23 收盤）", "€", None, 145.10, None, "Verified", "stockanalysis.com/quote/ams/BESI：2026-10-23 收盤 €145.10（−7.73%）", "marketscreener：前一日（10-22）收盤 €157.25，與 stockanalysis 歷史頁一致；10-23 收盤只有一個來源")
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重：情境 {s}", "%", None, w, None, "Assumed", "Andy 可改（AI 半導體 Project 共用設定 0.3／0.5／0.2）", "—")
inp("cons_rev26", "共識營收 FY2026", "€M", None, 1002, None, "Interested-party", "marketscreener 共識（頁面價格日 2026-10-09）", "stockanalysis：21 位分析師平均約 €1.0B（2026-10-23）")
inp("cons_eps26", "共識 EPS FY2026", "€", None, 4.334, None, "Interested-party", "marketscreener 共識（2026-10-09）", "stockanalysis 4.42（21 位，non-GAAP）")
inp("cons_rev27", "共識營收 FY2027", "€M", None, 1345, None, "Interested-party", "marketscreener 共識（2026-10-09，第三季財報前）", "第二來源：未找到")
inp("cons_eps27", "共識 EPS FY2027", "€", None, 6.448, None, "Interested-party", "marketscreener 共識（2026-10-09）", "第二來源：未找到")
inp("cons_rev28", "共識營收 FY2028", "€M", None, 1649, None, "Interested-party", "marketscreener 共識（2026-10-09）", "第二來源：未找到")
inp("cons_eps28", "共識 EPS FY2028", "€", None, 8.483, None, "Interested-party", "marketscreener 共識（2026-10-09）", "第二來源：未找到")
inp("lt_lo", "公司長期營收目標 下緣", "€M", None, 1700, None, "Interested-party", "投資人日 2026-06-19（quartr 摘要）：長期營收 €1.7–2.2B（前次 €1.5–1.9B）、營業利益率 45–55%；目標年度未載明", "2025 年版目標 €1.5–1.9B（Reuters／Yahoo 2025-06-12）")
inp("lt_hi", "公司長期營收目標 上緣", "€M", None, 2200, None, "Interested-party", "同上", "—")
inp("hb_cum25", "混合鍵合累計訂單 2025 年底", "台", None, 150, None, "Verified", "公司 Q4-25 新聞稿：18 家客戶、累計 150＋台", "—")
inp("hb_cum26q3", "混合鍵合累計訂單 2026 年第三季", "台", None, 215, None, "Interested-party", "investing.com Q3-26 簡報摘要：23 家客戶、累計 215＋台", "gurufocus 法說摘要同")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 60, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（百萬歐元；EPS 為歐元）。會計年度＝日曆年。藍字＝輸入；黑字＝公式（Derived）。季資料取自公司新聞稿（GlobeNewswire 轉載），年資料另以 stockanalysis 核對"
Q["A2"] = "FY2026 第三季（2026-07～09）已於 2026-10-22 公布並納入（J 欄）；第四季與全年預定 2027-02-25 公布（公司財務行事曆）。公布後更新本頁與 In 頁 q4g、q4gm、q4opx"
hdr(Q, 3, ["項目", "FY2024", "25Q1", "25Q2", "25Q3", "25Q4", "FY2025", "26Q1", "26Q2", "26Q3", "標記", "來源1", "來源2／說明"])
QR = {}
def qrow(r, key, label, vals, tag, s1, s2, fmt="#,##0.0"):
    QR[key] = r
    Q.cell(r, 1, label)
    for j, v in zip("BCDEFGHIJ", vals):
        if v is None: continue
        c = Q[f"{j}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    Q.cell(r, 11, tag); Q.cell(r, 12, s1); Q.cell(r, 13, s2)
qrow(4, "rev", "營收", [607.5, 144.1, 148.1, 132.7, 166.4, 591.3, 184.9, 249.9, 283.4], "Verified",
     "公司新聞稿：Q4-25／FY2025（2026-02-19）、Q2-26（2026-07-23）、Q3-26（2026-10-22）", "stockanalysis 年度：FY2024 607.47、FY2025 591.33；investing.com Q3-26 簡報摘要 283.4")
qrow(5, "gp", "毛利", [395.9, 91.7, 93.7, 82.6, 106.2, 374.3, 117.4, 164.1, 184.8], "Verified（26Q3 Derived）",
     "公司新聞稿；26Q3＝營收 283.4 × 毛利率 65.2%（簡報摘要）", "stockanalysis：FY2025 374.27、FY2024 395.94")
qrow(6, "gm", "毛利率", [f"={c}5/{c}4" for c in "BCDEFGHIJ"], "Derived", "—", "公司：FY2025 63.3%、26Q2 65.7%、26Q3 65.2%", "0.0%")
qrow(7, "opx", "營業費用（銷管＋研發，毛利 − 營業利益）", [f"={c}5-{c}8" for c in "BCDEFGHIJ"], "Derived", "—", "公司：26Q2 55.3、26Q3 55.1")
qrow(8, "oi", "營業利益", [195.6, 39.3, 43.5, 34.1, 56.2, 173.1, 63.9, 108.8, 129.7], "Verified", "公司新聞稿", "stockanalysis：FY2025 173.06、FY2024 195.59")
qrow(9, "fin", "財務費用淨額", [7.1, 3.0, 5.7, 5.1, 5.3, 19.1, 5.2, 6.3, "=J8-J10/(1-0.133)"], "Verified（26Q3 Derived）", "公司新聞稿；26Q3＝營業利益 − 淨利 ÷（1 − 13.3%）", "—")
qrow(10, "ni", "淨利", [182.0, 31.5, 32.1, 25.3, 42.8, 131.6, 51.6, 89.0, 107.2], "Verified", "公司新聞稿（FY2024 含 €18.2M 稅務利益）", "stockanalysis：FY2025 131.64、FY2024 181.99")
qrow(11, "eps", "EPS（基本，歐元）", [2.31, 0.40, 0.40, 0.32, 0.54, 1.66, 0.65, 1.11, 1.33], "Verified", "公司新聞稿", "stockanalysis 稀釋 EPS：FY2025 1.66、FY2024 2.30", "0.00")
qrow(12, "ord", "訂單", [586.7, 131.9, 128.0, 174.7, 250.4, 685.0, 269.7, 292.9, 342.6], "Verified", "公司新聞稿", "近 12 月訂單 €1.2B（Q3-26）")
qrow(13, "ck", "檢查：四季合計 − 全年（營收、淨利；B、C 欄）", ["=SUM(C4:F4)-G4", "=SUM(C10:F10)-G10"], "—", "—", "—")
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 10), ("eps", 11)):
    defname(f"a25_{nm}", f"Q!$G${rr}")
    for q, c in (("q1", "H"), ("q2", "I"), ("q3", "J")):
        defname(f"{q}_{nm}", f"Q!${c}${rr}")
Q.column_dimensions["A"].width = 46
for c in "BCDEFGHIJ": Q.column_dimensions[c].width = 11
Q.column_dimensions["K"].width = 22; Q.column_dimensions["L"].width = 80; Q.column_dimensions["M"].width = 80
Q.freeze_panes = "B4"

# ---------------- 推導引擎（PL 三情境各一個區塊；Sens 每個案例一個區塊） ----------------
COL = {y: chr(ord("C") + y - 2025) for y in range(2025, 2031)}   # C＝2025 … H＝2030
HB = {"A": "hbp_{y}_hi", "B": "hbp_{y}", "C": "hbp_{y}_lo"}
DEFP = dict(scn="A", den="", stk="", beta="beta", dlt="dlt", gt="g_tail", hb=None, att=None, thr="thr", asp="asp", shm=None, shl=None,
            hbsh="hb_sh26", el="el", prem="gm_prem", gmadj=True, F="opex_fix", v="opex_var", fin="fin", tax="tax", rev26="rev26", gp26="gp26", opx26="opx26", ni26="ni26", eps26="eps26")

def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    hb = P["hb"] if P["hb"] is not None else HB[s]
    att = P["att"] or ("att" + s + "_{y}")
    shm = P["shm"] or f"shm_{s}"; shl = P["shl"] or f"shl_{s}"
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.0", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for y, f in fs.items():
            ws[f"{COL[y]}{r}"] = f; ws[f"{COL[y]}{r}"].number_format = fmt
        ws.cell(r, 9, note); R[key] = r; r += 1
    ws.cell(r, 1, "A. HBM 混合鍵合（日曆年；與韓美同一晶粒口徑）").font = BOLD; r += 1
    row("ship", "HBM 出貨位元", "EB", {**{y: f"=hs{s}_{y % 100}" for y in YRS}, 2030: f"=G{r}*(1+{P['gt']})"}, fmt="0.00", note="markets/memory 同字母情境；2030＝2029 ×（1＋g_tail）")
    row("den", "DRAM 晶粒容量", "Gb", {y: f"=den_{y % 100}{P['den']}" for y in range(2025, 2031)}, fmt="0.0")
    row("dies", "核心晶粒數＝位元 × 8000 ÷ 容量", "百萬顆", {y: f"={COL[y]}{R['ship']}*8000/{COL[y]}{R['den']}" for y in range(2025, 2031)}, fmt="#,##0", note="1 EB＝8×10⁹ Gb；每顆核心晶粒鍵合一次")
    row("stk", "平均堆疊層數", "層", {y: f"=stk_{y % 100}{P['stk']}" for y in range(2025, 2031)}, fmt="0.0")
    row("lf", "層數加權（每顆晶粒鍵合時間）", "x", {y: f"=1+{P['beta']}*({COL[y]}{R['stk']}-$C{R['stk']})/4" for y in range(2025, 2031)}, fmt="0.000")
    row("hb", "混合鍵合占比", "%", {y: "=" + hb.format(y=y % 100) for y in range(2025, 2031)}, fmt="0.0%", note={"A": "情境 A：Disco 模型高值", "B": "情境 B：Disco 模型基準值", "C": "情境 C：Disco 模型低值"}[s] if P["hb"] is None else "")
    row("cap", "混合鍵合產能需求（晶粒 × 層數加權 × 混合鍵合占比）", "百萬顆當量", {y: f"={COL[y]}{R['dies']}*{COL[y]}{R['lf']}*{COL[y]}{R['hb']}" for y in range(2025, 2031)}, fmt="#,##0")
    row("tools", "所需機台＝產能需求 ÷ 每台年產能", "台", {y: f"={COL[y]}{R['cap']}/{P['thr']}" for y in range(2025, 2031)}, fmt="#,##0")
    row("newt", "新機需求＝MAX(0, 下一年 − 本年)＋汰換 × 本年", "台", {y: f"=MAX(0,{COL[y+1]}{R['tools']}-{COL[y]}{R['tools']})+{P['dlt']}*{COL[y]}{R['tools']}" for y in YRS}, fmt="#,##0", note="機台在產能需要前一年購入")
    ws.cell(r, 1, "B. 邏輯混合鍵合（SoIC 類；日曆年）").font = BOLD; r += 1
    um = lambda y: "1" if y < 2028 else f"um{y % 100}_{s}"
    row("cow", "AI 先進封裝晶圓需求＝台積電 CoWoS 需求 × 情境顆數係數", "千片", {**{y: f"=cow_{y % 100}*{um(y)}" for y in (2026, 2027, 2028, 2029)}, 2030: f"=G{r}*(1+{P['gt']})"}, fmt="#,##0", note="2030＝2029 ×（1＋g_tail）")
    row("att", "混合鍵合占 AI 先進封裝晶圓比例", "%", {y: "=" + att.format(y=y % 100) for y in range(2026, 2031)}, fmt="0.0%")
    row("idx", "混合鍵合邏輯晶圓（需求指數）", "千片當量", {y: f"={COL[y]}{R['cow']}*{COL[y]}{R['att']}" for y in range(2026, 2031)}, fmt="#,##0.0")
    row("newl", "新機需求指數＝MAX(0, 下一年 − 本年)＋汰換 × 本年", "千片當量", {y: f"=MAX(0,{COL[y+1]}{R['idx']}-{COL[y]}{R['idx']})+{P['dlt']}*{COL[y]}{R['idx']}" for y in range(2026, 2030)}, fmt="#,##0.0")
    ws.cell(r, 1, "C. 損益（會計年度；FY2026＝錨點）").font = BOLD; r += 1
    row("hbm", "HBM 混合鍵合營收＝新機 × 單價 × BESI 份額", "€M", {2025: "=0", 2026: "=hbm26", **{y: f"={COL[y]}{R['newt']}*{P['asp']}*{shm}" for y in (2027, 2028, 2029)}}, note="FY2025–26 尚未量產（Assumed 0）")
    row("lgc", "邏輯與其他混合鍵合營收", "€M", {2025: "=hb25", 2026: f"={P['rev26']}*{P['hbsh']}-hbm26", **{y: f"=$D{r + 1}*{COL[y]}{R['newl']}*{shl}" for y in (2027, 2028, 2029)}}, note="FY2026＝錨點營收 × 混合鍵合占比；FY2027 起＝k × 新機需求指數 × 相對份額")
    row("k", "校準係數 k＝FY2026 邏輯混合鍵合營收 ÷ FY2026 新機需求指數", "€M／千片當量", {2026: f"=D{R['lgc']}/D{R['newl']}"}, fmt="0.000", note="只用 FY2026（台積電 CoWoS 需求自 2026 起）；不使用共識")
    row("hbr", "混合鍵合營收合計", "€M", {y: f"={COL[y]}{R['hbm']}+{COL[y]}{R['lgc']}" for y in YRS})
    row("main", "主流業務營收（晶粒貼合、封裝、電鍍、TC 鍵合）", "€M", {2025: "=a25_rev-hb25", 2026: f"={P['rev26']}-D{R['hbr']}", **{y: f"={COL[y-1]}{r}*(1+{P['el']}*eqg{s}_{y % 100})" for y in (2027, 2028, 2029)}}, note="FY2027 起＝前一年 ×（1＋彈性 × 設備總額年增，markets/equipment 同字母）")
    row("rev", "營收", "€M", {y: f"={COL[y]}{R['hbr']}+{COL[y]}{R['main']}" for y in YRS})
    row("hsh", "混合鍵合占營收", "%", {y: f"={COL[y]}{R['hbr']}/{COL[y]}{R['rev']}" for y in YRS}, fmt="0.0%")
    adj = {2027: "0", 2028: "gm_adj_C" if (s == "C" and P["gmadj"]) else "0", 2029: "gm_adj_C" if (s == "C" and P["gmadj"]) else "0"}
    row("gm", "毛利率", "%", {2025: "=a25_gp/a25_rev", 2026: f"={P['gp26']}/{P['rev26']}", **{y: f"=$D{r}+{P['prem']}*({COL[y]}{R['hsh']}-$D{R['hsh']})+{adj[y]}" for y in (2027, 2028, 2029)}}, fmt="0.0%", note="FY2026＝錨點；FY2027 起＝FY2026＋溢價 × 混合鍵合占比變化（情境 C 另扣價格競爭）")
    row("gp", "毛利", "€M", {2025: "=a25_gp", 2026: f"={P['gp26']}", **{y: f"={COL[y]}{R['rev']}*{COL[y]}{R['gm']}" for y in (2027, 2028, 2029)}})
    row("opx", "營業費用", "€M", {2025: "=a25_opx", 2026: f"={P['opx26']}", **{y: f"={P['F']}+{P['v']}*{COL[y]}{R['rev']}" for y in (2027, 2028, 2029)}}, note="固定＋變動 × 營收")
    row("oi", "營業利益", "€M", {y: f"={COL[y]}{R['gp']}-{COL[y]}{R['opx']}" for y in YRS})
    row("oim", "營業利益率", "%", {y: f"={COL[y]}{R['oi']}/{COL[y]}{R['rev']}" for y in YRS}, fmt="0.0%")
    row("ni", "淨利", "€M", {2025: "=a25_ni", 2026: f"={P['ni26']}", **{y: f"=({COL[y]}{R['oi']}-{P['fin']})*(1-{P['tax']})" for y in (2027, 2028, 2029)}})
    row("eps", "EPS", "€", {2025: "=a25_eps", 2026: f"={P['eps26']}", **{y: f"={COL[y]}{R['ni']}/shares" for y in (2027, 2028, 2029)}}, fmt="0.00")
    row("units", "混合鍵合營收隱含機台數＝營收 ÷ 單價（對帳用）", "台", {y: f"={COL[y]}{R['hbr']}/{P['asp']}" for y in YRS}, fmt="#,##0")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — BESI 損益（€M；FY2025 實績，FY2026 錨點＝前三季實績＋第四季指引，FY2027–FY2029 為模型）。三情境各一個完整推導區塊（見下方），本頁頂端為摘要"
PB = {}; r0 = 40
TITLES = {"A": "情境 A（順利）：markets/memory、台積電顆數係數、markets/equipment 皆取 A；HBM 混合鍵合取 Disco 高值；邏輯導入快；份額守住",
          "B": "情境 B（平平）：上游皆取 B（2028 顆數 ×0.8、2029 ×0.85；設備 2028 見頂）；HBM 混合鍵合取 Disco 基準；份額小降",
          "C": "情境 C（不如預期）：上游皆取 C（2028–29 顆數 ×0.6；設備 2027 見頂、2028 −20%）；HBM 混合鍵合取 Disco 低值；邏輯導入慢、份額流失、毛利率下調"}
for s in SC:
    PB[s] = engine(PL, r0, TITLES[s], scn=s); r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "€M", "#,##0.0"), ("hbr", "  混合鍵合", "€M", "#,##0.0"), ("main", "  主流業務", "€M", "#,##0.0"),
                                ("gm", "毛利率", "%", "0.0%"), ("oi", "營業利益", "€M", "#,##0.0"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利", "€M", "#,##0.0"), ("eps", "EPS", "€", "0.00")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in YRS:
            PL[f"{COL[y]}{sr}"] = f"={COL[y]}{PB[s][key]}"; PL[f"{COL[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對：FY2026 三情境相同（營收、EPS 差額合計，應為 0）").font = BOLD
PL[f"C{sr}"] = f"=ABS(D{SUMR[('A','rev')]}-D{SUMR[('B','rev')]})+ABS(D{SUMR[('A','rev')]}-D{SUMR[('C','rev')]})+ABS(D{SUMR[('A','eps')]}-D{SUMR[('C','eps')]})"
PL.column_dimensions["A"].width = 56; PL.column_dimensions["B"].width = 14; PL.column_dimensions["I"].width = 50
for c in "CDEFGH": PL.column_dimensions[c].width = 12
PL["H3"] = "說明"; PL["H3"].font = BOLD; PL["H3"].fill = HDR
PL["H5"] = "FY2026 為錨點（前三季實績＋第四季指引），三情境相同；情境差異從 FY2027 起"

# ---------------- Units ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — HBM 晶粒 → 混合鍵合產能 → 機台；AI 先進封裝晶圓 → 邏輯混合鍵合需求指數（日曆年；三情境，引用 PL 區塊）"
hdr(UN, 3, ["項目", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "說明"])
ur = 4; UR = {}
for s in SC:
    UN.cell(ur, 1, f"情境 {s}").font = BOLD; ur += 1
    for key, lab, unit, fmt, y0, y1 in (("ship", "HBM 出貨位元", "EB", "0.00", 2025, 2031), ("dies", "核心晶粒數", "百萬顆", "#,##0", 2025, 2031),
                                        ("hb", "混合鍵合占比（HBM）", "%", "0.0%", 2025, 2031), ("tools", "HBM 混合鍵合所需機台", "台", "#,##0", 2025, 2031),
                                        ("newt", "HBM 混合鍵合新機需求", "台", "#,##0", 2025, 2030), ("cow", "AI 先進封裝晶圓需求", "千片", "#,##0", 2026, 2031),
                                        ("idx", "混合鍵合邏輯晶圓（需求指數）", "千片當量", "#,##0.0", 2026, 2031), ("newl", "邏輯新機需求指數", "千片當量", "#,##0.0", 2026, 2030),
                                        ("units", "混合鍵合營收隱含機台數", "台", "#,##0", 2025, 2030)):
        UN.cell(ur, 1, "  " + lab); UN.cell(ur, 2, unit)
        for y in range(y0, y1):
            UN[f"{COL[y]}{ur}"] = f"=PL!{COL[y]}{PB[s][key]}"; UN[f"{COL[y]}{ur}"].number_format = fmt
        UR[(s, key)] = ur; ur += 1
    ur += 1
UN.cell(ur, 9, "對照：混合鍵合累計訂單 2025 年底 150＋台、2026 年第三季 215＋台（公司）")
UN.column_dimensions["A"].width = 52; UN.column_dimensions["I"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值。跨公司規則：合理倍數＝10 年（2016–2025 年底）P/E 的分位，依循環位置：見頂年 25 分位、中段中位數、谷底 75 分位。循環位置依設備股共用規則（第 4 節）。本版未做先驗，無 Andy 合理 P/E 對照"
V["A2"] = "P/E 定義：日曆年底收盤價 ÷ 當年 EPS（本公司會計年度＝日曆年）。兩個來源各自取分位再平均（tools/valuation_rule.py 的多來源慣例）"
V.cell(4, 1, "0. 目標價與上檔空間").font = BOLD
hdr(V, 5, ["項目", "數值", "", "", "說明"])
VAL = {}
V.cell(6, 1, "目標價（12 個月，情境加權）"); V.cell(7, 1, "上檔空間（目標價 ÷ 現價 − 1）")
V.cell(8, 1, "情境 A 隱含價（順利）"); V.cell(9, 1, "情境 B 隱含價（平平）"); V.cell(10, 1, "情境 C 隱含價（不如預期）")
V.cell(11, 1, "情境權重 A／B／C")
for j, s in enumerate(SC): V.cell(11, 2 + j, f"=tp_w{s}").number_format = "0%"
V.cell(11, 5, "In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）")
V.cell(12, 1, "目標價方法")
V.cell(12, 2, "P/E：各情境 FY2027 EPS（模型）× 依循環位置的十年（2016–2025 年底）P/E 分位（設備股共用規則：A、B 中段＝中位數，C 見頂＝25 分位），再以 0.3／0.5／0.2 加權。偏誤：2016–2022 年（混合鍵合前）倍數 8–27 倍，2023–25 年 54–80 倍（盈餘下滑、混合鍵合溢價），分位受兩段樣本拉扯")

V.cell(14, 1, "1. 十年 P/E（年底）— 兩個來源").font = BOLD
hdr(V, 15, ["年底", "來源1：macrotrends（BESIY）", "來源2：stockanalysis（年底市值 ÷ 淨利）", "兩者差距", "標記", "來源2 出處"])
MT = {2016: 12.14, 2017: 12.33, 2018: 8.01, 2019: 27.37, 2020: 26.94, 2021: 18.62, 2022: 18.21, 2023: 60.96, 2024: 53.84, 2025: None}
S2 = {2016: (None, "未找到（stockanalysis 只到 FY2021）"), 2017: (None, "未找到"), 2018: (None, "未找到"), 2019: (None, "未找到"), 2020: (None, "未找到"),
      2021: (20.63, "stockanalysis 比率頁"), 2022: (18.73, "stockanalysis"), 2023: (59.45, "stockanalysis"), 2024: (57.95, "stockanalysis"), 2025: (80.10, "stockanalysis（來源1 2025 年底未找到）")}
for i, y in enumerate(range(2016, 2026), 16):
    V.cell(i, 1, f"{y}-12")
    if MT[y] is not None: V.cell(i, 2, MT[y]).font = BLUE
    if S2[y][0] is not None: V.cell(i, 3, S2[y][0]).font = BLUE
    V.cell(i, 4, f'=IF(OR(B{i}="",C{i}=""),"",B{i}-C{i})')
    for c in "BCD": V[f"{c}{i}"].number_format = "0.0"
    V.cell(i, 5, "Interested-party（資料商）"); V.cell(i, 6, S2[y][1])
V.cell(26, 1, "說明：macrotrends 以美國 ADR（BESIY）年底股價 ÷ 近 12 月 EPS（美元）；2025 年底值未找到（頁面最近一筆為 2025-06-17 的 60.7）。stockanalysis 以年底市值 ÷ 淨利（歐元），只有 2021–2025。兩源各自取分位再平均")
V.cell(28, 1, "2. 分位").font = BOLD
hdr(V, 29, ["統計", "採用（兩來源平均）", "來源1", "來源2", "用途"])
for i, (nm, p, use) in enumerate((("25 分位", 0.25, "見頂年"), ("中位數", 0.5, "中段"), ("75 分位", 0.75, "谷底")), 30):
    V.cell(i, 1, nm); V.cell(i, 5, use)
    V.cell(i, 3, f"=PERCENTILE(B16:B25,{p})"); V.cell(i, 4, f"=PERCENTILE(C16:C25,{p})"); V.cell(i, 2, f"=ROUND(AVERAGE(C{i},D{i}),2)")
    for c in "BCD": V[f"{c}{i}"].number_format = "0.0"
V.cell(33, 1, "年數"); V["C33"] = "=COUNT(B16:B25)"; V["D33"] = "=COUNT(C16:C25)"
defname("pe_peak", "Val!$B$30"); defname("pe_mid", "Val!$B$31"); defname("pe_trough", "Val!$B$32")

V.cell(35, 1, "4. 循環位置（設備股共用規則；見頂年度讀設備市場模型，不在本模型另訂）").font = BOLD
V.cell(36, 1, CP.rule_text(PEAK) + "。本公司會計年度＝日曆年，FY(t) 直接對照日曆年 t")
hdr(V, 37, ["情境", "設備市場見頂年度", "FY2027 對照日曆年", "判定", "倍數", "規則"])
for i, s in enumerate(SC, 38):
    V.cell(i, 1, s); V.cell(i, 2, f"=pk_{s}"); V.cell(i, 3, 2027).font = BLUE
    V.cell(i, 4, "=" + CP.excel_rule(f"C{i}", f"pk_{s}"))
    V.cell(i, 5, f'=IF(D{i}="見頂",pe_peak,IF(D{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 6, "tools/cycle_position.py")
V.cell(42, 1, "5. 隱含股價（歐元）：規則值＝FY2027 EPS × 依循環位置的倍數；其餘倍數僅供參考").font = BOLD
hdr(V, 43, ["情境", "EPS 年度", "EPS", "倍數類型", "倍數", "隱含股價", "相對現價", "備註"])
vr = 44
for s in SC:
    for y, typ, nm in ((2027, f"E{38 + SC.index(s)}", "rule"), (2027, "pe_peak", "參考：25 分位"), (2027, "pe_mid", "參考：中位數"), (2027, "pe_trough", "參考：75 分位"),
                       (2028, "pe_peak", "參考：25 分位"), (2028, "pe_mid", "參考：中位數"), (2028, "pe_trough", "參考：75 分位")):
        V.cell(vr, 1, s); V.cell(vr, 2, f"FY{y}"); V.cell(vr, 3, f"=PL!{COL[y]}{SUMR[(s,'eps')]}").number_format = "0.00"
        V.cell(vr, 4, nm if nm != "rule" else f'="規則值："&D{38 + SC.index(s)}'); V.cell(vr, 5, f"={typ}").number_format = "0.0"
        V.cell(vr, 6, f"=C{vr}*E{vr}").number_format = "#,##0.0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
        if nm == "rule":
            for c in range(1, 9): V.cell(vr, c).fill = KEY
            VAL[s] = vr
        vr += 1
for i, s in zip((8, 9, 10), SC):
    V.cell(i, 2, f"=F{VAL[s]}").number_format = "#,##0.0"; V.cell(i, 5, f"Val!F{VAL[s]}：FY2027 EPS × 規則倍數")
V.cell(6, 2, "=tp_wA*B8+tp_wB*B9+tp_wC*B10").number_format = "#,##0.0"; V.cell(6, 5, "權重 × 三情境隱含價")
V.cell(7, 2, "=B6/px-1").number_format = "0.0%"; V.cell(7, 5, "現價見第 6 節")
for c in range(1, 6): V.cell(6, c).fill = KEY
V.cell(vr + 1, 1, "6. 目前倍數").font = BOLD
hdr(V, vr + 2, ["項目", "A", "B", "C", "備註"])
fr = vr + 3
V.cell(fr, 1, "現價（stockanalysis 2026-10-23 收盤）"); V.cell(fr, 2, "=px").number_format = "#,##0.00"; PXROW = fr; fr += 1
for y in (2026, 2027, 2028):
    V.cell(fr, 1, f"現價 ÷ FY{y} EPS（模型）")
    for j, s in enumerate(SC): V.cell(fr, 2 + j, f"=px/PL!{COL[y]}{SUMR[(s,'eps')]}").number_format = "0.0"
    fr += 1
V.cell(fr, 1, "近 12 月 P/E（stockanalysis，2026-10-23）"); V.cell(fr, 2, 39.84).font = BLUE; V.cell(fr, 5, "近 12 月 EPS 3.64（同站）"); fr += 1
V.cell(fr, 1, "現價 ÷ 共識 FY2027 EPS（marketscreener 6.448）"); V.cell(fr, 2, "=px/cons_eps27").number_format = "0.0"; fr += 1
V.cell(fr, 1, "現價隱含倍數 ÷ 規則倍數（情境 A，FY2027）"); V.cell(fr, 2, f"=B{fr-4}/E38").number_format = "0.00"; V.cell(fr, 5, ">1＝現價高於規則給的倍數")
V.column_dimensions["A"].width = 50
for c in "BCDEFG": V.column_dimensions[c].width = 16
V.column_dimensions["F"].width = 60

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A）：對 FY2027、FY2028 營收與 EPS 的影響。每個案例是一個完整的推導區塊（本頁下方，與 PL 同一引擎；校準係數 k 在各區塊內重算）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收 €M", "FY2027 EPS €", "相對基準", "FY2028 營收 €M", "FY2028 EPS €", "相對基準"])
cases = [("HBM 混合鍵合 Disco 基準", "2028 5%、2029 15%、2030 25%", dict(hb="hbp_{y}")), ("HBM 混合鍵合 Disco 低值", "2028 0、2029 5%、2030 10%", dict(hb="hbp_{y}_lo")),
         ("邏輯導入 低（×0.7）", "2028 起各年比例 × 0.7", dict(att="attA_{y}_lo")), ("邏輯導入 高（×1.3）", "", dict(att="attA_{y}_hi")),
         ("邏輯導入 取情境 B 路徑", "", dict(att="attB_{y}")),
         ("FY2026 混合鍵合占比 18%", "k 較低、主流業務基期較高", dict(hbsh="hb_sh26_lo")), ("FY2026 混合鍵合占比 35%", "k 較高、主流業務基期較低", dict(hbsh="hb_sh26_hi")),
         ("機台有效產出 600／時", "機台數較多", dict(thr="(uph_lo*8760*util/1000000)")), ("機台有效產出 1,500／時", "", dict(thr="(uph_hi*8760*util/1000000)")),
         ("HBM 機台單價 €2.0M", "", dict(asp="asp_lo")), ("HBM 機台單價 €3.5M", "", dict(asp="asp_hi")),
         ("BESI HBM 份額 −10pt", "50%", dict(shm="shm_A_lo")), ("BESI HBM 份額 +10pt", "70%", dict(shm="shm_A_hi")),
         ("邏輯相對份額 0.9", "", dict(shl="shl_A_lo")),
         ("晶粒容量 低", "各年取低值（晶粒較多）", dict(den="_lo")), ("晶粒容量 高", "各年取高值（晶粒較少）", dict(den="_hi")),
         ("汰換比例 5%", "", dict(dlt="dlt_lo")), ("汰換比例 15%", "", dict(dlt="dlt_hi")),
         ("主流業務彈性 0.5", "", dict(el="el_lo")), ("主流業務彈性 1.5", "", dict(el="el_hi")),
         ("混合鍵合毛利溢價 0", "", dict(prem="gm_prem_lo")), ("混合鍵合毛利溢價 15pt", "", dict(prem="gm_prem_hi")),
         ("營業費用變動比例 25%", "", dict(F="(1-phi_lo)*opx26", v="phi_lo*opx26/rev26")), ("營業費用變動比例 60%", "", dict(F="(1-phi_hi)*opx26", v="phi_hi*opx26/rev26")),
         ("第四季營收 −5%", "錨點較低 → k、主流基期較低", dict(rev26="(q1_rev+q2_rev+q3_rev+q3_rev*(1+q4g_lo))")), ("第四季營收 +5%", "", dict(rev26="(q1_rev+q2_rev+q3_rev+q3_rev*(1+q4g_hi))")),
         ("稅率 16%", "", dict(tax="tax_hi"))]
def put_case(i, name, setting, R, pre=""):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    SE[f"C{i}"] = f"={pre}E{R['rev']}"; SE[f"D{i}"] = f"={pre}E{R['eps']}"; SE[f"F{i}"] = f"={pre}F{R['rev']}"; SE[f"G{i}"] = f"={pre}F{R['eps']}"
    SE[f"E{i}"] = f"=D{i}/D$4-1"; SE[f"H{i}"] = f"=G{i}/G$4-1"
    for c in "CF": SE[f"{c}{i}"].number_format = "#,##0"
    for c in "DG": SE[f"{c}{i}"].number_format = "0.00"
    for c in "EH": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
SB = {}
for n, (nm, setting, ov) in enumerate(cases):
    SB[nm] = engine(SE, 60 + n * 45, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    put_case(5 + n, nm, setting, SB[nm])
nn = 5 + len(cases)
put_case(nn, "情境 B", "上游皆取 B；HBM 混合鍵合基準；邏輯導入、份額取 B", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "上游皆取 C；HBM 混合鍵合低值；邏輯導入慢、份額流失、毛利率下調", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：機台比產能早一年買，FY2027 的營收看 CY2027→2028 的需求增量。邏輯段的校準係數 k 由 FY2026 營收推得，導入比例的水準被 k 吸收，只有比例的變化（斜率）影響之後各年；FY2026 混合鍵合占比同時改變 k 與主流業務基期，兩者方向相反。")
SE.column_dimensions["A"].width = 40; SE.column_dimensions["B"].width = 40
for c in "CDEFGH": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識、公司長期目標、累計訂單。只作事後對照，不為對上數字而調整參數。差距 > 10% 的原因見 README"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0", pt=False):
    global rr_
    RC.cell(rr_, 1, label); RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        L = chr(64 + cv)
        RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y=2026: (lambda s: f"=PL!{COL[y]}{SUMR[(s, key)]}")
RC.cell(rr_, 1, "A. FY2026 vs 共識（校準未使用）").font = BOLD; rr_ += 1
rec("營收", "=cons_rev26", pl("rev"), "marketscreener 共識（2026-10-09）", "錨點＝前三季實績＋第四季指引中點")
rec("EPS（歐元）", "=cons_eps26", pl("eps"), "marketscreener 共識（2026-10-09）", "stockanalysis 4.42", fmt="0.00")
RC.cell(rr_, 1, "B. FY2027–FY2028 vs 共識（第三季財報前的共識）").font = BOLD; rr_ += 1
rec("FY2027 營收", "=cons_rev27", pl("rev", 2027), "marketscreener")
rec("FY2027 EPS（歐元）", "=cons_eps27", pl("eps", 2027), "marketscreener", "", fmt="0.00")
rec("FY2028 營收", "=cons_rev28", pl("rev", 2028), "marketscreener")
rec("FY2028 EPS（歐元）", "=cons_eps28", pl("eps", 2028), "marketscreener", "", fmt="0.00")
RC.cell(rr_, 1, "C. 公司長期目標與累計訂單").font = BOLD; rr_ += 1
rec("FY2029 營收 vs 長期目標下緣 €1.7B", "=lt_lo", pl("rev", 2029), "投資人日 2026-06-19（目標年度未載明）", "Interested-party")
rec("FY2029 營收 vs 長期目標上緣 €2.2B", "=lt_hi", pl("rev", 2029), "同上")
rec("FY2026 混合鍵合隱含機台數 vs 2026 年前三季新增訂單（215 − 150）", "=hb_cum26q3-hb_cum25", lambda s: f"=PL!D{PB[s]['units']}", "公司：累計 150＋→ 215＋台", "口徑：營收（出貨）vs 訂單；邏輯機台單價可能與 HBM 不同")
RC.column_dimensions["A"].width = 60
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司新聞稿經 GlobeNewswire／stocktitan 轉載；法說與簡報取自摘要頁，原文待人工複核）"; SO["A1"].font = BOLD
srcs = [("公司：Q3-26 與前九月實績（2026-10-22；investing.com 簡報摘要）", "https://www.investing.com/news/company-news/besi-q3-2026-slides-revenue-doubles-net-margin-hits-38-93CH-4966911"),
 ("gurufocus：Q3-26 法說重點（混合鍵合 23 家客戶、215＋台、策略檢討、回購 €100M）", "https://www.gurufocus.com/news/9137691/be-semiconductor-industries-nv-besiy-q3-2026-earnings-call-highlights-revenue-more-than-doubles-as-hybrid-bonding-momentum-builds"),
 ("公司：Q2-26 與上半年實績（2026-07-23；stocktitan 轉載全文）", "https://www.stocktitan.net/news/BESIY/be-semiconductor-industries-n-v-announces-q2-26-and-h1-26-v0syj60xvlwf.html"),
 ("公司：Q2-26 新聞稿（besi.com）", "https://www.besi.com/investor-relations/press-releases/details/be-semiconductor-industries-nv-announces-q2-26-and-h1-26-results/"),
 ("公司：Q4-25 與 2025 全年實績（2026-02-19；stocktitan 轉載全文）", "https://www.stocktitan.net/news/BESIY/be-semiconductor-industries-n-v-announces-q4-25-and-full-year-2025-1hedov9aoxsi.html"),
 ("公司：財務行事曆（2027-02-25 公布 2026 全年）", "https://www.besi.com/investor-relations/financial-calendar/"),
 ("quartr：投資人日 2026 摘要（長期營收 €1.7–2.2B、營業利益率 45–55%）", "https://quartr.com/events/be-semiconductor-industries-n-v-besi-investor-day-2026_3eFihPox"),
 ("公司產品頁：Datacon 8800 CHAMEO ultra plus AC（2,000 CPH 乾循環、100nm）", "https://www.besi.com/products-technology/product-details/product/datacon-8800-chameo-ultra-plus-ac/"),
 ("tspasemiconductor：混合鍵合營收（FT 引 2023 年 €36M）", "https://tspasemiconductor.substack.com/p/hybrid-bonding-at-scale-besis-vision"),
 ("Morningstar：混合鍵合營收自 2023 年三倍（2025-02-21）", "https://morningstar.co.uk/uk/news/261234/be-semiconductor-guidance-disappoints;-long-term-outlook-strong.aspx"),
 ("stockanalysis：股價、統計、共識目標價、歷史收盤", "https://stockanalysis.com/quote/ams/BESI/forecast/"),
 ("stockanalysis：年度損益、年底 P/E（2021–2025）、ROE", "https://stockanalysis.com/quote/ams/BESI/financials/ratios/"),
 ("marketscreener：共識營收、淨利、EPS（2026–2028）", "https://in.marketscreener.com/quote/stock/BE-SEMICONDUCTOR-INDUSTRI-6318/finances/"),
 ("macrotrends：年底 P/E（2016–2024，ADR）、淨利率（2016–2025）", "https://www.macrotrends.net/stocks/charts/BESIY/be-semiconductor-industries/pe-ratio"),
 ("stockanalysis：同業前瞻 P/E（ASM International、ASMPT、Kulicke & Soffa、Disco、韓美）", "https://stockanalysis.com/quote/ams/ASM/statistics/"),
 ("市場層與上游成品：skhynix／disco／tsmc／markets/equipment 的 dist（本 repo）", ""),
 ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 80; SO.column_dimensions["B"].width = 120

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("BESI（BE Semiconductor Industries，BESI.AS）收支與估值模型 v0.1.1（AI 半導體 Project 步驟 5：設備與材料）　2026-10-25", True),
 ("v0.1（2026-10-25）：首版。樣板 hanmi/（鍵合設備，HBM 晶粒與混合鍵合比例口徑相同）與 advantest/（設備股統一循環規則）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("模型命題：BESI 的盈餘由混合鍵合機的新機需求決定——邏輯（SoIC 類）由 AI 先進封裝晶圓 × 混合鍵合導入比例推得，HBM 由 HBM 晶粒 × 混合鍵合占比（與韓美同一口徑）÷ 機台產能推得；傳統貼合機等主流業務隨設備市場循環。股價取決於 2027 年 CoWoS 增速放緩後，邏輯導入比例與 HBM4E／HBM5 混合鍵合能否接棒，以及 BESI 能否守住份額。", False), ("", False),
 ("會計年度＝日曆年。金額以百萬歐元（€M）為主；EPS 為歐元。FY2026 為錨點：前三季實績（184.9＋249.9＋283.4）＋第四季公司指引（營收持平 ±5%、毛利率 63–65%、費用 +5～10%）；第三季已於 2026-10-22 公布並納入。", False), ("", False),
 ("驅動因子 → 推導量", True),
 ("HBM：HBM 出貨位元（markets/memory 同字母情境）÷ 晶粒容量 → 核心晶粒數 × 層數加權 × 混合鍵合占比（Disco 模型，韓美用 1 − 此比例）→ 混合鍵合產能 ÷ 每台年產能 → 所需機台；新機＝下一年增量＋汰換；× 單價 × BESI 份額", False),
 ("邏輯：台積電 CoWoS 需求（基準）× 情境顆數係數（同字母）× 混合鍵合導入比例 → 需求指數；新機需求指數 × 校準係數 k（FY2026 混合鍵合營收）× 相對份額 → 邏輯混合鍵合營收；主流業務＝前一年 ×（1＋彈性 × markets/equipment 同字母設備總額年增）", False),
 ("毛利率：FY2026 錨點，之後加上混合鍵合占比變化 × 毛利溢價（C 扣價格競爭）；營業費用＝固定＋變動 × 營收；淨利＝（營業利益 − 財務費用）×（1 − 稅率）", False),
 ("情境 A／B／C＝順利／平平／不如預期：上游全部讀同字母（markets/memory HBM 位元、台積電情境顆數係數 2028 ×1.0／0.8／0.6、2029 ×1.0／0.85／0.6、markets/equipment 設備總額）；公司特有變數 A 最有利、C 最不利（HBM 混合鍵合 Disco 高／基準／低、邏輯導入快／中／慢、份額、C 另扣毛利率）。Val：十年 P/E 分位，循環位置依設備股共用規則（見頂年 " + "、".join(f"{s_} {PEAK[s_]}" for s_ in SC) + " → FY2027：" + CP.fy2027_text(PEAK) + "）；舊法規則值三情境 0.3／0.5／0.2 加權", False), ("", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司不揭露混合鍵合營收、產品別營收、機台單價與產能，相關輸入皆為 Assumed。", False),
 ("已知限制：(1) 混合鍵合營收、HBM／邏輯拆分、機台單價與產出皆未揭露（Assumed），FY2026 混合鍵合占比 25% 決定 k 與主流基期；(2) 邏輯導入比例無資料，水準被 k 吸收，只有斜率有作用；(3) 以台積電 CoWoS 需求代表 AI 邏輯先進封裝規模（Analogy），未含英特爾、消費性 CPU、共同封裝光學；(4) HBM 混合鍵合比例與晶粒、層數沿用 Disco 模型 Assumed，與韓美同進同退（方向相反）；(5) 機台假設比產能早一年購入；(6) 十年 P/E 2016–2020 只有一個來源、2025 只有一個來源；(7) 每年 €100M 回購與策略檢討（可能與應用材料合作）未建模。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

NIM = [0.1735, 0.2930, 0.2597, 0.2281, 0.3044, 0.3766, 0.3320, 0.3051, 0.2979, 0.2227]
CFG_V2 = dict(cols=COL, scn={"A": "順利", "B": "平平", "C": "不如預期"},
    peers=[('ASM International', 'ASM.AS', 25.77, '2026-10-23', 'stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）'),
           ('ASMPT', '0522.HK', 30.26, '2026-09-23', 'stockanalysis.com statistics：Forward PE（頁面價格日 2026-09-23，較舊）'),
           ('Kulicke & Soffa', 'KLIC', 12.61, '2026-10-23', 'stockanalysis.com statistics：Forward PE'),
           ('Disco', '6146.T', 26.01, '2026-10-23', 'stockanalysis.com statistics：Forward PE'),
           ('韓美半導體 Hanmi', '042700.KS', 44.63, '2026-08-07', 'stockanalysis.com statistics：Forward PE（頁面價格日 2026-08-07，較舊）')],
    fwd={2027: "2/12", 2028: "10/12"}, fwd_note='約 2027-10 時的未來 12 月 ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）',
    a3='目標價（DCF 與 Comparable 加權）：見本頁第 {row} 列起「0. 目標價」（放在既有內容之後）；原「0. 目標價與上檔空間」改名為「舊法：規則值情境加權（參考）」',
    cur='€', unit_rev='€M', unit_ni='€M', fmt='#,##0.00', fmt_rev='#,##0.0', band_fmt='#,##0.0',
    roe=('FY2023–FY2025', [0.3373, 0.3945, 0.2869], 'stockanalysis 比率頁（2026-10-25 讀取）：FY2023 33.73%、FY2024 39.45%、FY2025 28.69%', '單一來源；近 12 月 56.8%（同站，2026-10-23）'),
    nim=('2016–2025', NIM, 'macrotrends 淨利率（近 12 月，年底）：2016 17.35%、2017 29.30%、2018 25.97%、2019 22.81%、2020 30.44%、2021 37.66%、2022 33.20%、2023 30.51%、2024 29.79%、2025 22.27%', '2021–2025 與 stockanalysis 淨利 ÷ 營收相符（例：FY2025 131.64 ÷ 591.33＝22.3%）'),
    cons_tp=(270.57, 'stockanalysis（S&P Global）：23 位分析師平均 €270.57（中位數 255；區間 159–400），2026-10-23 頁面', 'marketscreener 共識目標價：未找到；高盛 2026-10-23 由 325 降至 277（marketscreener 新聞標題）'),
    cons_feps=("=2/12*cons_eps27+10/12*cons_eps28", '共識前瞻 EPS＝2／12 × 共識 FY2027 6.448＋10／12 × 共識 FY2028 8.483（marketscreener 2026-10-09，與本模型前瞻 EPS 同口徑）', 'Derived'),
    peer_note='後段與鍵合設備同業：ASM International（前段沉積，券商常與 BESI 並列）、ASMPT（TC／混合鍵合）、Kulicke & Soffa、Disco、韓美（TC 鍵合）；不含 A 股；韓華 Semitech、SEMES 未上市',
    dcf_note='；任一情境 DCF < 0 時取 0（股東有限責任），仍照 tp_w_dcf 加權',
    pe10='B16:B25',
    pe10_note='現價 ÷ 各情境 FY2028 EPS 在 2016–2024 年底 P/E（第 1 節來源1 macrotrends；2025 未找到）的分位；只作檢查，不進目標價',
    px_note='現價＝In 頁 px（2026-10-23 收盤）',
    method='DCF：股東現金流＝EPS ×（1 − g／ROE）（模型無折舊與資本支出項目；設備股輕資產），r 10%、g 4%、ROE 為近三年平均，FY2027–32 明確期（FY2030 起收斂法）＋永續，負值取 0。Comparable：ASM International、ASMPT、Kulicke & Soffa、Disco、韓美前瞻 P/E 中位數 × 12 個月後前瞻 EPS。偏誤：混合鍵合營收與單價未揭露（Assumed）；FY2027 受 CoWoS 增速放緩壓低，DCF 對導入比例斜率敏感；ASMPT、韓美的倍數日期較舊。',
    readme=['估值 v2（2026-10-25）：Val 頁「0. 目標價」＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，情境權重 A 0.3／B 0.5／C 0.2；負值取 0）＋0.5 × Comparable 目標價（同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解、歷史分位檢查。「舊法：規則值情境加權（參考）」不是目標價。',
            '預估期間延伸到 FY2032：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（PL 頁底部延伸區塊，灰底）。情境對應：A＝順利、B＝平平、C＝不如預期。本版未做先驗，無 Andy 合理 P/E 對照。'],
    nim_missing='')

# ---------------- 估值 v2：PL 延伸到 FY2032（收斂法）＋ Val「0. 目標價」（DCF 與 Comparable 加權） ----------------
def valuation_v2(CFG):
    cols = CFG["cols"]; MF = CFG["fmt"]; YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
    FADEFILL = PatternFill("solid", fgColor="E7E6E6")
    r = IN.max_row + 2
    IN.cell(r, 1, "J. 目標價（DCF 與 Comparable 加權）與 FY2030–32 收斂法參數（黃底＝Assumed，Andy 可改）").font = BOLD; r += 1
    def inrow(key, item, unit, base, tag, s1, s2, fmt=None, yellow=False):
        nonlocal r
        IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
        c = IN.cell(r, 5, base)
        if not (isinstance(base, str) and base.startswith("=")): c.font = BLUE
        if fmt: c.number_format = fmt
        if yellow: c.fill = KEY
        IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
        wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
        r += 1
    inrow("tp_r", "折現率（股東權益要求報酬；美、台、日、韓同一值）", "%", 0.10, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%）", "荷蘭公司沿用同一值", "0.0%", True)
    inrow("tp_g", "長期成長率（永續成長；收斂法 FY2032 營收年增的基準）", "%", 0.04, "Assumed", "Andy 可改（工作單預設 4%）", "情境 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt（PL 頁延伸區塊）", "0.0%", True)
    inrow("tp_roe", f"ROE（近三年平均，{CFG['roe'][0]}）", "%", "=AVERAGE(" + ",".join(str(x) for x in CFG["roe"][1]) + ")", "Derived", CFG["roe"][2], CFG["roe"][3], "0.0%")
    inrow("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", 0.5, "Assumed", "Andy 可改（2026-10-25 11:53：DCF／Comparable 各半）", "—", "0%", True)
    inrow("nim10", f"十年平均淨利率（{CFG['nim'][0]}，各年淨利 ÷ 營收的平均）", "%", "=AVERAGE(" + ",".join(str(x) for x in CFG["nim"][1]) + ")", "Derived", CFG["nim"][2], CFG["nim"][3], "0.0%")
    inrow("cons_tp", "市場共識目標價（賣方平均）", CFG["cur"], CFG["cons_tp"][0], "Interested-party", CFG["cons_tp"][1], CFG["cons_tp"][2], MF)
    inrow("cons_feps", "共識前瞻 EPS（差距拆解用）", CFG["cur"], CFG["cons_feps"][0], "Interested-party", CFG["cons_feps"][1], CFG["cons_feps"][2], MF)

    p0 = PL.max_row + 3
    PL.cell(p0, 1, "延伸到 FY2032：FY2025–FY2029 引用本頁頂端摘要；灰底 FY2030–FY2032 為收斂法，非驅動因子推導").font = BOLD
    PL.cell(p0 + 1, 1, "收斂法：營收年增由 FY2029 年增線性收斂到 FY2032 的長期成長（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt）；淨利率由 FY2029 線性收斂（A＝FY2029、B＝FY2029 與十年平均的中點、C＝十年平均 nim10）；股數維持 FY2029（EPS 與淨利同比例）")
    ext_cols = {y: chr(ord("C") + y - 2025) for y in range(2025, 2033)}
    hdr(PL, p0 + 2, ["情境（延伸）", "單位"] + [str(PL[f"{cols[y]}3"].value) for y in range(2025, 2030)] + ["FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    EXT = {}; pr = p0 + 3
    gadd = {"A": "+0.02", "B": "", "C": "-0.02"}
    for s in SC:
        PL.cell(pr, 1, f"情境 {s}（{CFG['scn'][s]}）").font = BOLD; pr += 1
        R = {k: pr + i for i, k in enumerate(("rev", "g", "nim", "ni", "eps"))}
        for k, lab in (("rev", "營收"), ("g", "  營收年增"), ("nim", "  淨利率（淨利 ÷ 營收）"), ("ni", "淨利（收斂法延伸）"), ("eps", "EPS（延伸）")):
            PL.cell(R[k], 1, lab); PL.cell(R[k], 2, {"rev": CFG["unit_rev"], "ni": CFG["unit_ni"], "eps": CFG["cur"]}.get(k, "%"))
        for y in range(2025, 2033):
            c = ext_cols[y]; pc = ext_cols.get(y - 1)
            if y <= 2029:
                PL[f"{c}{R['rev']}"] = f"={cols[y]}{SUMR[(s, 'rev')]}"
                PL[f"{c}{R['ni']}"] = f"={cols[y]}{SUMR[(s, 'ni')]}"
                PL[f"{c}{R['eps']}"] = f"={cols[y]}{SUMR[(s, 'eps')]}"
                PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
                if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            else:
                k = y - 2029
                mT = {"A": f"$G{R['nim']}", "B": f"($G{R['nim']}+nim10)/2", "C": "nim10"}[s]
                PL[f"{c}{R['g']}"] = f"=$G{R['g']}+((tp_g{gadd[s]})-$G{R['g']})*{k}/3"
                PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
                PL[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({mT}-$G{R['nim']})*{k}/3"
                PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
                PL[f"{c}{R['eps']}"] = f"=$G{R['eps']}*{c}{R['ni']}/$G{R['ni']}"
                for kk in R: PL[f"{c}{R[kk]}"].fill = FADEFILL
            for kk, fm in (("rev", CFG["fmt_rev"]), ("g", "0.0%"), ("nim", "0.0%"), ("ni", CFG["fmt_rev"]), ("eps", MF)):
                if PL[f"{c}{R[kk]}"].value is not None: PL[f"{c}{R[kk]}"].number_format = fm
        PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導")
        EXT[s] = R; pr += 6
    for c in "IJ": PL.column_dimensions[c].width = 14
    PL.column_dimensions["K"].width = 40

    REN = (("0. 目標價與上檔空間", "舊法：規則值情境加權（參考）"), ("目標價（12 個月，情境加權）", "規則值情境加權價（12 個月；參考）"),
           ("上檔空間（目標價 ÷ 現價 − 1）", "相對現價（規則值情境加權價 ÷ 現價 − 1）"), ("目標價方法", "規則值方法"))
    for row in V.iter_rows(min_col=1, max_col=1):
        v = row[0].value
        if isinstance(v, str):
            for a, b in REN:
                if v.startswith(a): row[0].value = b + v[len(a):]; break
    t0 = V.max_row + 3
    T = {}
    def tset(rr, key, label, fs, fmt, note=""):
        V.cell(rr, 1, label); T[key] = rr
        for c, f in zip("BCD", fs):
            if f is None: continue
            V[f"{c}{rr}"] = f; V[f"{c}{rr}"].number_format = fmt
        if note: V[f"E{rr}"] = note
    V.cell(t0, 1, "0. 目標價").font = BOLD; T["top"] = t0
    V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD; T["h01"] = t0 + 1
    hdr(V, t0 + 2, ["項目", f"A（{CFG['scn']['A']}）", f"B（{CFG['scn']['B']}）", f"C（{CFG['scn']['C']}）", "說明"]); T["hd1"] = t0 + 2
    V.cell(t0 + 3, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{t0 + 3}"] = "PL 頁延伸區塊；FY2030 起為收斂法"; T["eps_h"] = t0 + 3
    EPSR = {}
    for i, y in enumerate(YRS6):
        tset(t0 + 4 + i, f"eps{y}", f"  FY{y}", [f"=PL!{ext_cols[y]}{EXT[s]['eps']}" for s in SC], MF); EPSR[y] = t0 + 4 + i
    rr = t0 + 10
    V.cell(rr, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); V[f"E{rr}"] = "模型無折舊與資本支出項目，故不用（淨利＋折舊－資本支出）；設備股資本支出低（2026 第三季資本支出含資本化研發 €8.5M），此式偏誤小"; T["cf_h"] = rr
    tset(rr + 1, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
    RV = rr + 1; CFR = {}
    for i, y in enumerate(YRS6):
        tset(rr + 2 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], MF); CFR[y] = rr + 2 + i
    c1, c6 = CFR[2027], CFR[2032]
    tset(rr + 8, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], MF,
         "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027" + CFG["dcf_note"])
    tset(rr + 9, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{rr + 8}*(1+tp_r)" for c in "BCD"], MF)
    tset(rr + 10, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{rr + 9}+tp_wB*C{rr + 9}+tp_wC*D{rr + 9}", None, None], MF, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁 tp_wA／tp_wB／tp_wC")
    q = rr + 12
    V.cell(q, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; T["h02"] = q
    hdr(V, q + 1, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["phd"] = q + 1
    for i, (nm, tk, pe, dt, src) in enumerate(CFG["peers"], q + 2):
        V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
    pl_ = q + 1 + len(CFG["peers"]); T["p1"] = q + 2; T["pl"] = pl_
    q = pl_ + 1
    tset(q, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{T['p1']}:C{pl_})", None, None], "0.0", CFG["peer_note"])
    fw = CFG["fwd"]
    def feps(c): return "(" + "+".join(f"{w}*{c}{EPSR[y]}" for y, w in fw.items()) + ")"
    tset(q + 1, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=tp_wA*" + feps("B") + "+tp_wB*" + feps("C") + "+tp_wC*" + feps("D"), None, None], MF, CFG["fwd_note"])
    tset(q + 2, "comptp", "Comparable 目標價", [f"=B{q}*B{q + 1}", None, None], MF, "同業中位數 × 上列 EPS")
    q += 4
    V.cell(q, 1, "0.3 結論").font = BOLD; T["h03"] = q
    hdr(V, q + 1, ["項目", "數值", "", "", "說明"]); T["hd3"] = q + 1
    q += 2
    tset(q, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], MF, "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁")
    tset(q + 1, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], MF, "引用 0.2")
    tset(q + 2, "cons", "市場共識目標價", ["=cons_tp", None, None], MF, CFG["cons_tp"][1])
    tset(q + 3, "upw", "上檔空間：加權目標價", [f"=B{q}/px-1", None, None], "0.0%", CFG["px_note"])
    tset(q + 4, "upc", "上檔空間：Comparable 目標價", [f"=B{q + 1}/px-1", None, None], "0.0%")
    tset(q + 5, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型 12 個月後前瞻 EPS ÷ 共識前瞻 EPS − 1；" + CFG["cons_feps"][1])
    tset(q + 6, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%")
    V[f"E{q + 6}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
    tset(q + 7, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"{CFG["band_fmt"]}")&"～"&TEXT(B{T["dcf12"]},"{CFG["band_fmt"]}")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
    PE = CFG["pe10"]
    tset(q + 8, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
         [f"=IF(px/{c}{EPSR[2028]}<MIN({PE}),0,IF(px/{c}{EPSR[2028]}>MAX({PE}),1,PERCENTRANK({PE},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%", CFG["pe10_note"])
    V.cell(q + 9, 1, "目標價方法"); V[f"B{q + 9}"] = CFG["method"]; T["meth"] = q + 9
    for k in ("dcftp", "comptp", "tpw"): V[f"B{T[k]}"].fill = KEY
    if CFG.get("a3"): V["A3"] = CFG["a3"].format(row=t0)

    s0 = SE.max_row + 3
    SE.cell(s0, 1, "目標價參數（只改 DCF 參數與權重，PL 收斂法的長期成長不隨 tp_g 變動；DCF 負值取 0）").font = BOLD
    hdr(SE, s0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
    def dcf_formula(c, r_="tp_r", g="tp_g"):
        cf = lambda y: f"Val!{c}{EPSR[y]}*(1-MIN({g}/tp_roe,0.8))"
        pv = "+".join(f"{cf(y)}/(1+{r_})^{t}" for t, y in enumerate(YRS6, 1))
        return f"MAX(0,{pv}+{cf(2032)}*(1+{g})/({r_}-{g})/(1+{r_})^6)"
    cases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r_="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r_="(tp_r-0.01)")),
             ("長期成長 +1pt", "g 5%（永續與再投資率）", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
             ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
    for k, (nm, st, kw) in enumerate(cases):
        i = s0 + 2 + k; w = kw.pop("w", "tp_w_dcf")
        SE.cell(i, 1, nm); SE.cell(i, 2, st)
        rr_ = kw.get("r_", "tp_r")
        SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, **kw)}" for s, c in zip(SC, "BCD")) + f")*(1+{rr_})"
        SE[f"D{i}"] = f"=Val!B{T['comptp']}"
        SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
        SE[f"F{i}"] = f"=E{i}/E${s0 + 2}-1"
        for c in "CDE": SE[f"{c}{i}"].number_format = MF
        SE[f"F{i}"].number_format = "0.0%"
    T["sens0"] = s0; T["sens_last"] = s0 + 1 + len(cases)
    rr = RD.max_row + 1
    for t in CFG["readme"]:
        RD.cell(rr, 1, t); rr += 1
    return T, EXT

V2T, V2EXT = valuation_v2(CFG_V2)

# ---------------- v0.1.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00、14:17）；新增列放在各頁現有最後一列之後，既有儲存格不變 ----------------
_U = f"B{V2T['upw']}"
_G = V.max_row + 2
V.cell(_G, 1, "上檔空間級距")
V[f"B{_G}"] = f'=IF({_U}>0.5,"大於 50%",IF({_U}>0.25,"25%–50%",IF({_U}>=0,"0%–25%",IF({_U}>=-0.25,"0%～−25%",IF({_U}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"B{_G}"].number_format = "@"; V[f"E{_G}"] = "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）"
V.cell(_G + 1, 1, "投資建議")
V[f"B{_G + 1}"] = f'=IF({_U}>0.5,"強力買進",IF({_U}>0.25,"買進",IF({_U}>=-0.25,"中立",IF({_U}>=-0.5,"放空","強力放空"))))'
V[f"B{_G + 1}"].number_format = "@"
V[f"E{_G + 1}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空"
OBS = ["目標價 €133.69（DCF 97.68、Comparable 169.71 各半），比現價 145.10（2026-10-23）低 7.9%，級距 0%～−25%；Comparable 單獨為 +17.0%。DCF 低是情境分歧造成：DCF 情境價值 A 158.4（現價的 1.09 倍）、B 82.4、C 44.8，B 的 FY2028 EPS 6.59 比 A 的 9.37 低 30%。",
       "本模型 12 個月後前瞻 EPS 6.52 比共識前瞻 EPS 8.14 低 19.9%；共識目標價 270.57 隱含 33.2 倍，同業中位數 26.0 倍（倍數因素 +27.7%）。",
       "最敏感的是邏輯（SoIC 類）混合鍵合導入速度：導入比例 ×0.7／×1.3 使 FY2027 EPS −22.0%／+22.0%，改走情境 B 路徑 −12.2%；HBM 混合鍵合取 Disco 低值 −7.2%。下一個檢驗事件是 Q4-26 財報（約 2027-02 下旬）的混合鍵合訂單（Q3-26 訂單 €342.6M）。"]   # 重點觀察（與 besi/README.md 相同文字）
_R = RD.max_row + 1
RD.cell(_R, 1, "v0.1.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁最後新增「上檔空間級距」「投資建議」兩列（依「上檔空間：加權目標價」固定規則）；README 頁最後新增投資建議與重點觀察三點。")
RD.cell(_R + 1, 1, f'="投資建議："&Val!B{_G + 1}&"（依上檔空間固定規則）"')
for _i, _t in enumerate(OBS, 1): RD.cell(_R + 1 + _i, 1, f"重點觀察 {_i}：{_t}")
if os.environ.get("G1_ROWS"): print("G1ROWS", "besi", "grade", _G, "rec", _G + 1, "readme_obs", _R + 2, _R + 4)

if os.environ.get("V2_ROWS"):
    print("V2", V2T, V2EXT)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# 相對倍數調整（Andy 2026-10-27「照建議做」；tools/relative_multiple.py，2026-10-29 接入）
import os, sys  # noqa: E401
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools"))
import relative_multiple as RMX  # noqa: E402
RMX.apply(wb, "besi")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tools", "in_layout")); import apply_in_layout  # noqa: E402
apply_in_layout.apply(wb, "v0.1.2", "2026-10-31", split=True, split_skip=[])  # v0.1.2：In 頁版面；v0.1.3（2026-10-31）：第 3 點情境專屬列拆分（數值不變）
wb.save(sys.argv[1])
print("PB", {s: PB[s] for s in SC}); print("SUMR", SUMR); print("VAL", VAL, "PXROW", PXROW, "fwd rows to", fr, "SENS_LAST", SENS_LAST, "Recon last", rr_)
print("UR", UR)
