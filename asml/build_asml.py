# ASML（ASML.AS／ASML ADR）收支與估值模型 v0.2（AI 半導體 Project 步驟 5：設備與材料）
# 用法：python3 asml/build_asml.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年（FY2026＝2026-01～12）。金額：歐元（Q 頁 €M，PL 頁 €B）；EPS 為歐元。
# 結構比照 advantest/build_advantest.py：README、In、Q、Units、PL、Val、Recon、Sens、Sources。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
YRS = (2025, 2026, 2027, 2028, 2029)

# ---------------- 讀入設備市場模型成品（唯讀；不修改） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
EQ_F = only("markets/equipment/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ws_ = openpyxl.load_workbook(EQ_F, data_only=True)["Mkt"]
SER = {"la": "先進邏輯（台積電設備＋其他代工中屬先進者）〔AI〕", "lo": "其他代工與邏輯", "dr": "DRAM（含 HBM）〔AI〕", "na": "NAND", "wf": "晶圓前段設備"}
EQ = {}; cnt = {k: 0 for k in SER}
for r in ws_.iter_rows(min_row=119):
    lab = r[0].value
    for k, nm in SER.items():
        if lab == nm and cnt[k] < 3:
            s = SC[cnt[k]]; cnt[k] += 1
            EQ[(k, s)] = ([r[2 + j].value for j in range(5)], f"{rel(EQ_F)}｜Mkt!C{r[0].row}:G{r[0].row}")
assert all(v == 3 for v in cnt.values()), cnt
for r in ws_.iter_rows(min_row=40, max_row=50):
    if r[0].value == "SEMI 晶圓前段設備（WFE）":
        SEMI_WFE = ([r[2 + j].value for j in range(5)], f"{rel(EQ_F)}｜Mkt!C{r[0].row}:G{r[0].row}")
# 2025–2027 三情境相同（設備市場模型的設計）；建置時檢查
for k in SER:
    for j in range(3):
        assert abs(EQ[(k, "A")][0][j] - EQ[(k, "C")][0][j]) < 1e-9, (k, j)
PEAK, PEAK_SRC = CP.peak_years(ROOT)

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：Q 頁與指引為百萬歐元（€M），驅動與 PL 為十億歐元（€B）；設備市場為十億美元（US$B）。每個輸入附兩個來源（來源2 找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司數字取自 SEC 6-K（ASML 新聞稿、財務報表、簡報）與法說逐字稿；asml.com 的 PDF 經搜尋引擎摘錄讀取"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)

sec("A. 晶圓前段設備支出（markets/equipment 成品，三情境；唯讀，不自行重估）")
SNM = {"la": "先進邏輯設備", "lo": "其他代工與邏輯設備", "dr": "DRAM（含 HBM）設備", "na": "NAND 設備", "wf": "晶圓前段設備合計"}
for k in SER:
    for s in SC:
        for j, y in enumerate(YRS):
            inp(f"{k}{s}_{y % 100}", f"{SNM[k]} {y}（情境 {s}）", "US$B", None, EQ[(k, s)][0][j], None, "Derived",
                EQ[(k, s)][1] if j == 0 else "同上", "設備市場模型：公司資本支出 × 設備占比＋其他廠；2025–2027 三情境相同，2028–29 依情境係數" if (j == 0 and s == "A") else "同上")
for j, y in enumerate(YRS):
    inp(f"semi_wfe_{y % 100}", f"SEMI 晶圓前段設備（WFE）{y}", "US$B", None, SEMI_WFE[0][j], None, "Interested-party" if j else "Verified", SEMI_WFE[1] if j == 0 else "同上",
        "SEMI 2025 實績；2026 起為 SEMI 預測（設備商協會，傾向樂觀；只用於對帳）" if j == 0 else "同上")
for s in SC:
    inp(f"pk_{s}", f"情境 {s}：設備市場見頂年度（日曆年；循環位置規則用）", "年", None, PEAK[s], None, "Derived", PEAK_SRC[s],
        "設備股共用規則（tools/cycle_position.py）；ASML 會計年度＝日曆年，FY(t) 直接對照日曆年 t" if s == "A" else "同上")

sec("B. FY2026 第四季公司指引（2026-10-14）與全年指引；FY2026 錨點＝前三季實績＋第四季指引中點")
inp("g_q4_rev", "第四季營收指引（14.2–15.2 中點）", "€M", 14200, 14700, 15200, "Interested-party", "ASML Q3 2026 新聞稿（SEC 6-K，2026-10-14）：Q4 €14.2–15.2B", "簡報（6-K 附件）同；stockhub.kr 摘要同")
inp("g_q4_ibm", "第四季裝機基礎管理（IBM）營收指引", "€M", None, 3700, None, "Interested-party", "ASML Q3 2026 簡報：Q4 IBM 約 €3.7B", "法說逐字稿（Benzinga）同")
inp("g_q4_gm", "第四季毛利率指引", "%", 0.56, 0.57, 0.58, "Interested-party", "ASML Q3 2026 新聞稿：約 57%", "stockhub.kr 摘要同；低／高 ±1pt 為 Assumed")
inp("g_q4_rd", "第四季研發費用指引", "€M", None, 1200, None, "Interested-party", "同上：約 €1.2B", "同上")
inp("g_q4_sga", "第四季銷管費用指引", "€M", None, 400, None, "Interested-party", "同上：約 €0.4B", "同上")
inp("g_fy_rev", "FY2026 營收指引（44–45 中點）", "€M", 44000, 44500, 45000, "Interested-party", "ASML Q3 2026 新聞稿：€44–45B（7 月 €43–45B、4 月 €36–40B、1 月 €34–39B）", "aiweekly 2026-10-15：下緣上調至 €44B")
inp("g_fy_gm", "FY2026 毛利率指引", "%", None, 0.56, None, "Interested-party", "同上：約 56%（7 月 54–56%、1 月 51–53%）", "法說逐字稿（Benzinga）同")
inp("g_tax", "FY2026 有效稅率指引", "%", None, 0.17, None, "Interested-party", "Q3 2026 簡報：約 17%（Q3 實際 16.1%）", "法說逐字稿同")
inp("g_euv", "FY2026 EUV 系統營收成長（指引「超過 45%」）", "%", None, 0.45, None, "Interested-party", "Q3 2026 法說：EUV 營收成長超過 45%；Low-NA EUV 出貨約 65 台（2025：44 台）", "Q2 2026 法說同（超過 45%、約 65 台）")
inp("g_log", "FY2026 邏輯系統營收成長（指引）", "%", None, 0.10, None, "Interested-party", "Q3 2026 法說：邏輯約 +10%、記憶體約 +75%", "Q2 2026 法說：記憶體超過 +75%、先進邏輯代工超過 +25%（口徑不同）")
inp("g_mem", "FY2026 記憶體系統營收成長（指引）", "%", None, 0.75, None, "Interested-party", "同上", "同上")
inp("g_ibm26", "FY2026 IBM 成長（指引）", "%", None, 0.50, None, "Interested-party", "Q3 2026 法說：約 +50%（7 月：超過 +30%）", "錨點值見 Q 頁（由前三季＋第四季指引推得，對照用）")
inp("sh_q4", "第四季加權平均股數（基本）", "百萬股", 380.0, 381.5, 383.1, "Assumed", "無：公司未給；Q3 2026 加權 383.1M（6-K），Q3 買回約 2.8M 股", "高＝Q3 值（不再買回）；基準＝Q3 再減約 0.4%")

sec("C. 系統營收的組成（FY2026 錨點的拆分）")
inp("a25_log", "FY2025 邏輯系統營收", "€M", None, 16100, None, "Verified", "ASML Q4 2025 法說逐字稿（2026-01-28）：邏輯 €16.1B（+22%）", "記憶體 €8.4B；兩者合計 24.5 與 6-K 系統營收 24,474 相符")
inp("a25_mem", "FY2025 記憶體系統營收", "€M", None, 8400, None, "Verified", "同上：記憶體 €8.4B（−2%）", "同上")
inp("log_sh26", "FY2026 邏輯占系統營收（指引成長推得）", "%", None, "=a25_log*(1+g_log)/(a25_log*(1+g_log)+a25_mem*(1+g_mem))", None, "Derived", "16.1×1.10 ÷（16.1×1.10＋8.4×1.75）", "Q1–Q3 2026 單季：49%／51%／54%（法說）")
inp("cn_sh26", "FY2026 中國占系統營收", "%", 0.17, 0.20, 0.23, "Interested-party", "公司：2026 年中國約占總營收 20%（Q4 2025、Q3 2026 法說）；本模型套用於系統營收（Assumed）", "aiweekly 2026-10-15：Q3 占系統出貨 23%、Q2 14%；低／高＝17%／23%（Assumed）")
inp("cn_tot25", "FY2025 中國占總營收", "%", None, 0.33, None, "Interested-party", "Caixin 2026-01-29（公司 FY2025 揭露）：33%，高於公司 2025-07 預估的約 25%", "2024 約 36%（Caixin 時間軸，AI 生成，弱）；第二來源：未找到")

sec("D. 驅動：設備支出指數 × 微影強度與價格 → 系統需求；產能上限；中國；IBM")
inp("g_int", "微影強度與價格提升（每年；相對晶圓前段設備支出）", "%/年", 0.0, 0.03, 0.06, "Assumed", "無直接實證。公司：關鍵產品漲價自 2027 生效、2028 全年生效（Q3 2026 法說）；EUV 層數增加", "Analogy：ASML 系統營收（美元）÷ SEMI WFE 2025＝23.7%（Recon）；只有一年可比，不作校準")
inp("cap_27", "2027 產能成長（Low-NA EUV 與 DUV 浸潤式）", "%", 0.20, 0.30, 0.35, "Interested-party", "Q2、Q3 2026 法說：2027 產能約 +30%", "Q1 2026 法說：Low-NA EUV 2027「至少 80 台」（2026 約 65 台 → +23%）；低＝20%（Assumed）")
inp("cap_28", "2028 產能成長", "%", 0.0, 0.15, 0.30, "Assumed", "公司：研究 2028 再 +30%（未確定）", "基準取一半；高＝公司研究值")
inp("cap_29", "2029 產能成長", "%", 0.0, 0.10, 0.20, "Assumed", "無", "—")
for s, g, lo, hi in (("A", 0.0, -0.10, 0.10), ("B", -0.10, -0.20, 0.0), ("C", -0.25, -0.40, -0.10)):
    inp(f"g_cn_{s}", f"中國系統營收年成長 FY2027–29（情境 {s}）", "%", lo, g, hi, "Assumed",
        "Analogy：中國占總營收 2024 約 36%→2025 33%→2026 約 20%（公司）；CEO：2027 以後中國需求「仍非常強」（Q3 2026 法說）" if s == "A" else ("出口管制擴大、成熟製程擴產放緩" if s == "B" else "管制擴及浸潤式與維修服務（Assumed 下行）"),
        "第二來源：未找到")
for s, g in (("A", 0.12), ("B", 0.08), ("C", 0.03)):
    inp(f"g_ibm_{s}", f"IBM 年成長 FY2027–29（情境 {s}）", "%", g - 0.04, g, g + 0.04, "Assumed",
        "Analogy：IBM 2024 €6,494M→2025 €8,193M（+26%）→2026 錨點約 +49%（公司：EUV 裝機量與生產力升級）" if s == "A" else ("升級需求隨設備支出持平而降溫" if s == "B" else "設備下行時升級需求回落，服務仍隨裝機量成長"),
        "第二來源：未找到")

sec("E. 毛利率與營業費用（兩點推導：FY2025 實績與 FY2026 錨點；不寫死利益率）")
inp("gm_b", "毛利率對營收規模的斜率（每 ln 單位營收）", "pt", "=gm_b*0.5", "=(a26_gp/a26_rev-a25_gp/a25_rev)/LN(a26_rev/a25_rev)", "=gm_b*1.5", "Derived",
    "兩點：FY2025 52.8%（€32.7B）、FY2026 錨點（€44.5B）；毛利率＝FY2026 錨點毛利率＋斜率 × ln(營收 ÷ FY2026 營收)", "公司 2030 目標：營收 €44–60B 時毛利率 56–60%（2024 投資人日；Recon 對照）。低／高＝×0.5／×1.5（Assumed）")
inp("opex25", "營業費用 FY2025（研發＋銷管）", "€M", None, "=a25_rd+a25_sga", None, "Derived", "Q 頁", "—")
inp("opex26", "營業費用 FY2026 錨點（前三季＋第四季指引）", "€M", None, "=a26_rd+a26_sga", None, "Derived", "Q 頁", "—")
inp("phi", "營業費用中隨營收變動的比例（以 FY2026 為基準）", "%", 0.0, "=(opex26-opex25)/(a26_rev-a25_rev)*a26_rev/opex26", 0.40, "Derived",
    "基準＝兩點斜率 × FY2026 營收 ÷ FY2026 營業費用；營業費用＝FY2026 ×［(1−phi)＋phi × 營收 ÷ FY2026 營收］", "FY2025→FY2026 營收 +36%、營業費用 +4%；低＝全固定、高 40%（Assumed）")
inp("nioi", "淨利 ÷ 營業利益（FY2027 起）", "x", 0.80, "=a25_ni/a25_oi", 0.87, "Derived", "基準＝FY2025 實績 0.850（稅率約 17%＋利息與權益法收益）", "FY2026 前三季含 Mistral AI 投資評價益 €509.7M（一次性），不當基準；低／高 Assumed")
inp("bb", "股數每年減少（買回）", "%/年", 0.0, 0.01, 0.02, "Assumed", "公司：€12B 買回計畫，2026–2028 執行（Q4 2025 簡報）", "以 €1,400／股約 8.6M 股、3 年約 −0.75%／年；Q3 2026 單季買回 2.8M 股（€4.1B），基準取 1%")

sec("F. EUV 顯示與產能核對（不影響 EPS）")
for y, v in ((2027, 0.55), (2028, 0.57), (2029, 0.58)):
    inp(f"euv_sh_{y % 100}", f"EUV 占系統營收 {y}", "%", v - 0.05, v, v + 0.05, "Assumed", "Analogy：FY2025 47%（11.6÷24.5）、FY2026 錨點約 52%", "第二來源：未找到")
for y, v in ((2027, 8), (2028, 12), (2029, 15)):
    inp(f"hna_n_{y % 100}", f"High-NA EUV 營收認列台數 {y}", "台", v - 4, v, v + 5, "Assumed", "公司：2026 前三季認列 4 台（Q1 2、Q2 1、Q3 1）；三星 2028、台積電 2030 量產（Q3 2026 法說）", "第二來源：未找到")
inp("hna_asp", "High-NA EUV 單價", "€M／台", 300, 350, 400, "Interested-party", "媒體：約 US$380M（Tom's Hardware／Yole 轉述）≈ €350M", "公司未揭露；第二來源：未找到")
inp("lna_asp", "Low-NA EUV 單價（NXE:3800E）", "€M／台", 200, 220, 250, "Assumed", "Analogy：FY2025 EUV €11.6B ÷ 48 台＝€242M（含 High-NA）", "公司未揭露；第二來源：未找到")
inp("euv_cap27", "2027 Low-NA EUV 產能", "台", None, 80, None, "Interested-party", "Q1 2026 法說：2027「至少 80 台」", "Q2、Q3 法說：約 +30%（65×1.3≈85）")

sec("G. 匯率、股價、共識與公司長期目標")
inp("fx_25", "歐元兌美元 FY2025 平均", "US$/€", None, 1.13, None, "Verified", "芬蘭央行（ECB 參考匯率）2025 年平均 1.1300", "只用於對帳（ASML 系統營收換成美元 ÷ 設備市場）")
inp("fx_f", "歐元兌美元 FY2026 以後（Assumed）", "US$/€", 1.05, 1.16, 1.25, "Assumed", "無：模型以指數連結設備市場，匯率不進營收；只用於對帳", "ECB 2025-12-31 參考匯率 1.1750")
inp("px", "股價（Euronext Amsterdam，2026-10-23 收盤）", "€", None, 1407.2, None, "Interested-party", "stockanalysis.com（AMS:ASML）2026-10-23 收盤 €1,407.20", "eulerpool 同為 €1,407.20")
inp("sa_fpe", "前瞻 P/E（stockanalysis）", "x", None, 24.48, None, "Interested-party", "stockanalysis.com 2026-10-23", "近 12 月 P/E 40.24、EPS 34.97（同站）")
inp("cons_rev26", "共識營收 FY2026", "€M", None, 44460, None, "Interested-party", "stockanalysis.com（35 位分析師；2026-10-17 更新）", "第二來源：未找到（eulerpool 只有 EPS）")
inp("cons_eps26", "共識 EPS FY2026", "€", None, 41.21, None, "Interested-party", "stockanalysis.com（non-GAAP 調整後）", "eulerpool：41.24")
inp("cons_rev27", "共識營收 FY2027", "€M", None, 57210, None, "Interested-party", "stockanalysis.com 摘要「Next Year」", "第二來源：未找到")
inp("cons_eps27", "共識 EPS FY2027", "€", None, 55.93, None, "Interested-party", "stockanalysis.com 摘要「Next Year」", "eulerpool：P/E 2027 24.81 → 隱含 €56.7（Derived）")
inp("eul_pe28", "eulerpool P/E 2028（共識）", "x", None, 18.80, None, "Interested-party", "eulerpool.com（ASML.AS）", "隱含 FY2028 EPS＝現價 ÷ 18.80；第二來源：未找到")
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")
inp("t30_rev_lo", "公司 2030 營收目標 下緣", "€M", None, 44000, None, "Interested-party", "2024 投資人日（Q4 2025 簡報重申）：約 €44–60B", "—")
inp("t30_rev_hi", "公司 2030 營收目標 上緣", "€M", None, 60000, None, "Interested-party", "同上", "下次更新：2027-06-10 投資人日")
inp("t30_gm_lo", "公司 2030 毛利率目標 下緣", "%", None, 0.56, None, "Interested-party", "同上：約 56–60%", "—")
inp("t30_gm_hi", "公司 2030 毛利率目標 上緣", "%", None, 0.60, None, "Interested-party", "同上", "—")

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for col, w in zip("ABCDEFGHI", [13, 56, 10, 9, 11, 9, 16, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績（€M） ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（百萬歐元；EPS 為歐元）。會計年度＝日曆年。藍字＝輸入（SEC 6-K 公司財報）；黑字＝公式（Derived）。FY2026 錨點＝前三季實績＋第四季公司指引（2026-10-14）"
Q["A2"] = "FY2026 第三季（2026-10-14 公布）已納入；第四季與全年 2026 財報預定 2027-01 下旬公布（公司慣例：1 月第四個週三；確切日期：未找到），屆時以實績取代 J 欄指引"
HQ = ["項目", "9M 2025（1–9月）", "Q3 2025", "H1 2025（Derived）", "Q4 2025", "FY2025", "Q1 2026", "Q2 2026", "Q3 2026", "Q4 2026（指引）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"]
hdr(Q, 3, HQ)
QR = {}
SRC_Q = "SEC 6-K：Q3 2026 財務報表（季度摘要表 Q3 2025–Q3 2026、9M 2025／2026）"
def qrow(r, key, label, vals, tag, s1, s2, fmt="#,##0.0"):
    QR[key] = r; Q.cell(r, 1, label)
    for j, v in zip("BCDEFGHIJK", vals):
        if v is None: continue
        c = Q[f"{j}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    Q.cell(r, 12, tag); Q.cell(r, 13, s1); Q.cell(r, 14, s2)
qrow(4, "sys", "系統營收（net system sales）", [16890.3, 5553.8, "=B4-C4", 7584.0, "=B4+E4", 6279.4, 6564.8, 8373.2, "=J6-J5", "=SUM(G4:J4)"], "Verified", SRC_Q, "Q4 2025 簡報：FY2025 €24,474M；Q3 2026 簡報：Q3 €8.4B")
qrow(5, "ibm", "裝機基礎管理（IBM；net service and field option sales）", [6058.9, 1962.2, "=B5-C5", 2134.1, "=B5+E5", 2487.5, 2761.7, 3288.8, "=g_q4_ibm", "=SUM(G5:J5)"], "Verified／Interested-party", SRC_Q + "；J 欄＝公司指引", "Q4 2025 簡報：FY2025 IBM €8.2B（+26%）")
qrow(6, "rev", "營收", ["=B4+B5", "=C4+C5", "=D4+D5", "=E4+E5", "=F4+F5", "=G4+G5", "=H4+H5", "=I4+I5", "=g_q4_rev", "=SUM(G6:J6)"], "Derived", "系統＋IBM；J 欄＝指引中點", "stockanalysis：FY2025 €32,667M；Q3 2026 新聞稿 €11,662M")
qrow(7, "gp", "毛利", [12189.5, 3880.3, "=B7-C7", 5068.6, "=B7+E7", 4645.0, 5035.4, 6699.8, "=J6*g_q4_gm", "=SUM(G7:J7)"], "Verified", SRC_Q, "stockanalysis：FY2025 毛利率 52.83%")
qrow(8, "gm", "毛利率", [f"={c}7/{c}6" for c in "BCDEFGHIJK"], "Derived", "—", "公司：FY2025 52.8%；Q3 2026 57.5%；FY2026 指引約 56%", "0.0%")
qrow(9, "rd", "研發費用", [3436.5, 1108.7, "=B9-C9", 1262.3, "=B9+E9", 1184.9, 1276.6, 1174.4, "=g_q4_rd", "=SUM(G9:J9)"], "Verified", SRC_Q, "—")
qrow(10, "sga", "銷管費用", [882.5, 303.2, "=B10-C10", 375.2, "=B10+E10", 302.3, 302.7, 352.6, "=g_q4_sga", "=SUM(G10:J10)"], "Verified", SRC_Q, "—")
qrow(11, "oi", "營業利益", [f"={c}7-{c}9-{c}10" for c in "BCDEFGHIJK"], "Derived", "毛利 − 研發 − 銷管", "公司：Q3 2026 營業利益 €5,172.8M（營業利益率 44.4%）；FY2025 營業利益率 34.6%")
qrow(12, "oir", "營業利益（公司公布；核對用）", [7870.5, 2468.4, "=B12-C12", 3431.1, "=B12+E12", 3157.8, 3456.1, 5172.8, None, None], "Verified", SRC_Q, "stockhub.kr：Q3 2026 營業利益 €5.2B")
qrow(13, "eqg", "權益投資評價益（Mistral AI；一次性）", [0, 0, "=B13-C13", 0, "=B13+E13", 0, 0, 509.7, 0, "=SUM(G13:J13)"], "Verified", SRC_Q + "：Gain on equity investments", "Q3 2026 法說：初始投資的評價上升約 €0.5B")
qrow(14, "ni", "淨利", [6769.8, 2124.5, "=B14-C14", 2839.6, "=B14+E14", 2756.7, 2917.6, 4960.4, "=J11*nioi", "=SUM(G14:J14)"], "Verified／Derived", SRC_Q + "；J 欄＝第四季營業利益 × 淨利÷營業利益（FY2025 實績比例）", "Q4 2025 簡報：FY2025 €9,609M")
qrow(15, "shr", "加權平均股數（基本，百萬股）", [389.2, 387.3, None, 386.5, "=F14/F16", 385.4, 384.5, 383.1, "=sh_q4", "=AVERAGE(G15:J15)"], "Verified／Derived", SRC_Q + "；FY2025＝淨利 ÷ 基本 EPS", "J 欄為 Assumed（In 頁 sh_q4）", "#,##0.0")
qrow(16, "eps", "EPS（基本，歐元）", [17.39, 5.49, "=B16-C16", 7.35, 24.73, 7.15, 7.59, 12.95, "=J14/J15", "=K14/K15"], "Verified／Derived", SRC_Q + "；FY2025 基本 24.73（Q4 2025 簡報）", "稀釋 EPS：FY2025 24.71、Q3 2026 12.93；FY2026 錨點＝淨利 ÷ 四季平均股數", "0.00")
qrow(17, "units", "微影系統銷售台數（新＋二手）", [225, 72, "=B17-C17", 102, "=B17+E17", 79, 91, 121, None, None], "Verified", SRC_Q, "Q4 2025 簡報：FY2025 327 台", "0")
qrow(18, "ck", "檢查：營業利益（自算 − 公布）", [f"={c}11-{c}12" for c in "BCDEFGHI"], "—", "—", "應為 0", "0.0")
Q["J4"].number_format = "#,##0.0"
for nm, rr in (("sys", 4), ("ibm", 5), ("rev", 6), ("gp", 7), ("rd", 9), ("sga", 10), ("oi", 11), ("ni", 14), ("shr", 15), ("eps", 16)):
    defname(f"a25_{nm}", f"Q!$F${rr}"); defname(f"a26_{nm}", f"Q!$K${rr}")

Q.cell(21, 1, "系統營收的組成（€M；公司法說）").font = BOLD
hdr(Q, 22, ["項目", "FY2025", "Q1 2026", "Q2 2026", "Q3 2026", "Q4 2026（Derived）", "FY2026 錨點", "", "", "", "", "標記", "來源1（公司）", "來源2／說明"])
segs = [("EUV 系統營收", [11600, 4100, 3800, 4100, "=G23-SUM(C23:E23)", "=B23*(1+g_euv)"], "Verified／Derived", "Q4 2025 簡報：FY2025 €11.6B（+39%，48 台 NXE＋EXE）；Q1 法說「超過 €4.1B」（含 2 台 High-NA）、Q2 €3.8B（1 台）、Q3 €4.1B（1 台）", "FY2026＝FY2025 × (1＋45%)（公司「超過 45%」，故為下限）；Q4＝全年 − 前三季"),
        ("非 EUV 系統營收（DUV＋量測檢測）", ["=Q!F4-B23", "=Q!G4-C23", "=Q!H4-D23", "=Q!I4-E23", "=Q!J4-F23", "=Q!K4-G23"], "Derived", "系統營收 − EUV", "公司：FY2025 DUV €12.0B（279 台，47% 浸潤式）、量測檢測 €825M；Q2 非 EUV €2.8B、Q3 €4.3B"),
        ("EUV 占系統營收", ["=B23/Q!F4", "=C23/Q!G4", "=D23/Q!H4", "=E23/Q!I4", "=F23/Q!J4", "=G23/Q!K4"], "Derived", "—", "—"),
        ("邏輯占系統營收", ["=a25_log/Q!F4", 0.49, 0.51, 0.54, None, "=log_sh26"], "Verified／Derived", "Q1、Q2、Q3 2026 法說：49%／51%／54%；FY2025＝€16.1B ÷ 24,474", "FY2026＝指引成長推得（In 頁 log_sh26）"),
        ("中國占比", ["=cn_tot25", None, 0.14, 0.23, None, "=cn_sh26"], "Interested-party", "FY2025 占總營收 33%（Caixin，公司揭露）；Q2、Q3 占系統出貨 14%／23%（aiweekly 2026-10-15）", "FY2026：公司約 20%（總營收），本模型套用於系統營收；Q1：未找到")]
for i, (lab, vals, tag, s1, s2) in enumerate(segs, 23):
    Q.cell(i, 1, lab)
    for j, v in zip("BCDEFG", vals):
        if v is None: continue
        c = Q[f"{j}{i}"]; c.value = v; c.number_format = "0.0%" if i >= 25 else "#,##0"
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    Q.cell(i, 12, tag); Q.cell(i, 13, s1); Q.cell(i, 14, s2)
defname("a25_euv", "Q!$B$23"); defname("a26_euv", "Q!$G$23")
Q.cell(29, 1, "讀法：FY2026 錨點營收 44,455 與指引中點 44,500 差 −0.1%；毛利率 55.7%（指引約 56%）。前三季淨利含 Mistral AI 評價益 509.7（一次性），FY2026 EPS 為報告口徑；FY2027 起淨利 ÷ 營業利益用 FY2025 實績比例（不含此項）。")
Q.column_dimensions["A"].width = 50
for c in "BCDEFGHIJK": Q.column_dimensions[c].width = 13
Q.column_dimensions["L"].width = 20; Q.column_dimensions["M"].width = 80; Q.column_dimensions["N"].width = 70
Q.freeze_panes = "B4"

# ---------------- Units：設備市場連結（三情境，US$B）與指數 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 晶圓前段設備支出（markets/equipment，US$B；三情境）→ 邏輯與記憶體指數（FY2026＝1）。ASML 會計年度＝日曆年，不需 3/4、1/4 轉換"
hdr(UN, 3, ["項目", "單位"] + [f"CY{y}" for y in YRS] + ["說明"])
YC = {y: chr(ord("C") + y - 2025) for y in YRS}    # C＝2025 … G＝2029
UR = {}; ur = 4
for s in SC:
    UN.cell(ur, 1, f"情境 {s}").font = BOLD; ur += 1
    for key, lab, f in (("lg", "邏輯晶圓前段設備（先進邏輯＋其他代工與邏輯）", lambda y: f"=la{s}_{y%100}+lo{s}_{y%100}"),
                        ("mm", "記憶體晶圓前段設備（DRAM 含 HBM＋NAND）", lambda y: f"=dr{s}_{y%100}+na{s}_{y%100}"),
                        ("wf", "晶圓前段設備合計", lambda y: f"=wf{s}_{y%100}")):
        UN.cell(ur, 1, "  " + lab); UN.cell(ur, 2, "US$B")
        for y in YRS: UN[f"{YC[y]}{ur}"] = f(y); UN[f"{YC[y]}{ur}"].number_format = "0.0"
        UR[(s, key)] = ur; ur += 1
    for key, lab, src in (("li", "  邏輯指數（2026＝1）", "lg"), ("mi", "  記憶體指數（2026＝1）", "mm")):
        UN.cell(ur, 1, lab); UN.cell(ur, 2, "x")
        for y in YRS: UN[f"{YC[y]}{ur}"] = f"={YC[y]}{UR[(s, src)]}/$D{UR[(s, src)]}"; UN[f"{YC[y]}{ur}"].number_format = "0.000"
        UR[(s, key)] = ur; ur += 1
    UN.cell(ur, 1, "  其他（類比、功率、光罩、廠務）＝合計 − 邏輯 − 記憶體"); UN.cell(ur, 2, "US$B")
    for y in YRS: UN[f"{YC[y]}{ur}"] = f"={YC[y]}{UR[(s,'wf')]}-{YC[y]}{UR[(s,'lg')]}-{YC[y]}{UR[(s,'mm')]}"; UN[f"{YC[y]}{ur}"].number_format = "0.0"
    UN.cell(ur, 8, "不進 ASML 的需求指數（其他前段的微影含量低；Assumed）"); ur += 2
UN.cell(ur, 1, "讀法：邏輯指數 2027 約 1.16、記憶體約 1.35（三情境相同）；2028–29 依設備市場情境係數。設備市場的其他代工與記憶體廠含中國廠（中芯、長鑫等）；本模型另以中國系統營收列處理中國，指數的中國成分重疊（限制，見 README）。")
UN.column_dimensions["A"].width = 60; UN.column_dimensions["B"].width = 8; UN.column_dimensions["H"].width = 70
for c in "CDEFG": UN.column_dimensions[c].width = 11

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
FYC = YC
DEFP = dict(gi="g_int", cn="cn_sh26", gcn=None, gibm=None, c28="cap_28", c27="cap_27", gmb="gm_b", phi="phi", nioi="nioi", bb="bb", scn="A")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    gcn = P["gcn"] or f"g_cn_{s}"; gibm = P["gibm"] or f"g_ibm_{s}"
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    F3 = (2027, 2028, 2029)
    ws.cell(r, 1, "A. 設備市場連結（Units 頁；情境別）").font = BOLD; r += 1
    row("li", "邏輯晶圓前段設備指數（2026＝1）", "x", {FYC[y]: f"=Units!{FYC[y]}{UR[(s,'li')]}" for y in YRS}, "0.000")
    row("mi", "記憶體晶圓前段設備指數（2026＝1）", "x", {FYC[y]: f"=Units!{FYC[y]}{UR[(s,'mi')]}" for y in YRS}, "0.000")
    row("up", "微影強度與價格係數（2026＝1）", "x", {FYC[y]: f"=(1+{P['gi']})^{y-2026}" for y in (2026,) + F3}, "0.000", "(1＋g_int)^(t−2026)")
    ws.cell(r, 1, "B. 系統營收（€B）").font = BOLD; r += 1
    row("nlog", "非中國 邏輯系統需求", "€B", {"D": f"=a26_sys/1000*(1-{P['cn']})*log_sh26", **{FYC[y]: f"=$D{r}*{FYC[y]}{R['li']}*{FYC[y]}{R['up']}" for y in F3}}, note="FY2026＝錨點系統營收 ×（1−中國占比）× 邏輯占比；之後 × 邏輯指數 × 強度係數")
    row("nmem", "非中國 記憶體系統需求", "€B", {"D": f"=a26_sys/1000*(1-{P['cn']})*(1-log_sh26)", **{FYC[y]: f"=$D{r}*{FYC[y]}{R['mi']}*{FYC[y]}{R['up']}" for y in F3}})
    row("cn", "中國系統需求", "€B", {"D": f"=a26_sys/1000*{P['cn']}", **{FYC[y]: f"={chr(ord(FYC[y])-1)}{r}*(1+{gcn})" for y in F3}}, note="FY2026＝錨點 × 中國占比；之後 × (1＋中國年成長)")
    row("dem", "系統需求合計", "€B", {FYC[y]: f"={FYC[y]}{R['nlog']}+{FYC[y]}{R['nmem']}+{FYC[y]}{R['cn']}" for y in (2026,) + F3})
    capf = {2027: f"(1+{P['c27']})", 2028: f"(1+{P['c27']})*(1+{P['c28']})", 2029: f"(1+{P['c27']})*(1+{P['c28']})*(1+cap_29)"}
    row("cap", "系統產能上限（FY2026 錨點 × 產能成長 × 強度係數）", "€B", {"D": "=a26_sys/1000", **{FYC[y]: f"=$D{r}*{capf[y]}*{FYC[y]}{R['up']}" for y in F3}}, note="公司稱 2026 產能已滿（供不應求）；上限只在需求超過時生效")
    row("sys", "系統營收＝MIN(需求, 產能)", "€B", {"C": "=a25_sys/1000", "D": "=a26_sys/1000", **{FYC[y]: f"=MIN({FYC[y]}{R['dem']},{FYC[y]}{R['cap']})" for y in F3}})
    row("bind", "產能限制生效（1＝是）", "", {FYC[y]: f"=IF({FYC[y]}{R['dem']}>{FYC[y]}{R['cap']},1,0)" for y in F3}, "0")
    row("cnsh", "中國占系統營收", "%", {FYC[y]: f"={FYC[y]}{R['cn']}/{FYC[y]}{R['dem']}" for y in (2026,) + F3}, "0.0%", "產能限制時各組成同比例縮減")
    row("ibm", "裝機基礎管理（IBM）營收", "€B", {"C": "=a25_ibm/1000", "D": "=a26_ibm/1000", **{FYC[y]: f"={chr(ord(FYC[y])-1)}{r}*(1+{gibm})" for y in F3}})
    row("rev", "營收", "€B", {FYC[y]: f"={FYC[y]}{R['sys']}+{FYC[y]}{R['ibm']}" for y in YRS})
    row("revg", "營收年增", "%", {FYC[y]: f"={FYC[y]}{R['rev']}/{chr(ord(FYC[y])-1)}{R['rev']}-1" for y in (2026,) + F3}, "0.0%")
    ws.cell(r, 1, "C. 損益").font = BOLD; r += 1
    row("gm", "毛利率", "%", {"C": "=a25_gp/a25_rev", "D": "=a26_gp/a26_rev", **{FYC[y]: f"=$D{r}+{P['gmb']}*LN({FYC[y]}{R['rev']}/$D{R['rev']})" for y in F3}}, "0.0%", "FY2026 錨點毛利率＋斜率 × ln(營收 ÷ FY2026 營收)")
    row("gp", "毛利", "€B", {FYC[y]: f"={FYC[y]}{R['rev']}*{FYC[y]}{R['gm']}" for y in YRS})
    row("opx", "營業費用（研發＋銷管）", "€B", {"C": "=opex25/1000", "D": "=opex26/1000", **{FYC[y]: f"=$D{r}*((1-{P['phi']})+{P['phi']}*{FYC[y]}{R['rev']}/$D{R['rev']})" for y in F3}}, note="FY2026 ×［(1−phi)＋phi × 營收比］")
    row("oi", "營業利益", "€B", {FYC[y]: f"={FYC[y]}{R['gp']}-{FYC[y]}{R['opx']}" for y in YRS})
    row("oim", "營業利益率", "%", {FYC[y]: f"={FYC[y]}{R['oi']}/{FYC[y]}{R['rev']}" for y in YRS}, "0.0%")
    row("ni", "淨利", "€B", {"C": "=a25_ni/1000", "D": "=a26_ni/1000", **{FYC[y]: f"={FYC[y]}{R['oi']}*{P['nioi']}" for y in F3}}, note="FY2026 含 Mistral AI 評價益 €0.51B（一次性）")
    row("sh", "加權平均股數", "百萬股", {"C": "=a25_shr", "D": "=a26_shr", **{FYC[y]: f"=sh_q4*(1-{P['bb']})^{y-2026}" for y in F3}}, "#,##0.0")
    row("eps", "EPS（基本）", "€", {"C": "=a25_eps", "D": "=a26_eps", **{FYC[y]: f"={FYC[y]}{R['ni']}*1000/{FYC[y]}{R['sh']}" for y in F3}})
    ws.cell(r, 1, "D. 顯示與核對（不影響 EPS）").font = BOLD; r += 1
    row("euv", "EUV 系統營收", "€B", {"C": "=a25_euv/1000", "D": "=a26_euv/1000", **{FYC[y]: f"={FYC[y]}{R['sys']}*euv_sh_{y%100}" for y in F3}})
    row("lna", "隱含 Low-NA EUV 台數（EUV 營收 − High-NA）÷ 單價", "台", {FYC[y]: f"=({FYC[y]}{R['euv']}*1000-hna_n_{y%100}*hna_asp)/lna_asp" for y in F3}, "0")
    row("usd", "系統營收（美元）", "US$B", {"C": f"=C{R['sys']}*fx_25", **{FYC[y]: f"={FYC[y]}{R['sys']}*fx_f" for y in (2026,) + F3}}, "0.0")
    row("lsh", "ASML 系統營收 ÷ 晶圓前段設備（設備市場模型）", "%", {FYC[y]: f"={FYC[y]}{R['usd']}/Units!{FYC[y]}{UR[(s,'wf')]}" for y in YRS}, "0.0%", "微影（ASML）占晶圓前段設備；FY2026 起匯率為 Assumed")
    R["end"] = r
    return R

# ---------------- PL：三情境 ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — ASML 損益（€B；EPS 為歐元）。FY2025 為實績，FY2026 為錨點（前三季實績＋第四季公司指引；三情境相同），FY2027–FY2029 為模型。三情境各一個完整推導區塊（見下方），本頁頂端為摘要"
SNAME = {"A": "情境 A：基準（設備支出續增；中國持平；IBM +12%）", "B": "情境 B：設備支出 2028 持平、2029 −5%；中國每年 −10%；IBM +8%",
         "C": "情境 C：設備支出 2028 −20%；中國每年 −25%（管制擴大）；IBM +3%"}
PB = {}; r0 = 50
for s in SC:
    PB[s] = engine(PL, r0, SNAME[s], scn=s); r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "€B", "#,##0.0"), ("sys", "  系統營收", "€B", "#,##0.0"), ("ibm", "  裝機基礎管理（IBM）", "€B", "#,##0.0"), ("cnsh", "  中國占系統營收", "%", "0.0%"),
        ("gm", "毛利率", "%", "0.0%"), ("oi", "營業利益", "€B", "#,##0.0"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利", "€B", "#,##0.0"), ("eps", "EPS", "€", "#,##0.00"))
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in YRS:
            if key == "cnsh" and y == 2025: continue
            PL[f"{FYC[y]}{sr}"] = f"={FYC[y]}{PB[s][key]}"; PL[f"{FYC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對：FY2026 錨點 EPS − Q 頁（三情境，應為 0）"); PL[f"D{sr}"] = "=" + "+".join(f"ABS(D{SUMR[(s,'eps')]}-a26_eps)" for s in SC); PL[f"D{sr}"].number_format = "0.0000"
PL.column_dimensions["A"].width = 52; PL.column_dimensions["B"].width = 10; PL.column_dimensions["H"].width = 60
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Val：P/E 估值 ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（歐元）。跨公司規則：合理倍數＝10 年（2016–2025 年底）P/E 的分位，依循環位置：見頂年 25 分位、中段中位數、谷底 75 分位。循環位置依設備股共用規則（第 4 節；ASML 會計年度＝日曆年）"
V["A2"] = "P/E 定義（自算）：年底 ADR 收盤價（美元，未調整）÷（當年基本 EPS（歐元）× ECB 年底歐元兌美元）。ADR 與阿姆斯特丹普通股 1：1；當年即年底所在會計年度（會計年度＝日曆年）"
V.cell(14, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 15, ["項目", "A", "B", "C", "說明"])
V.cell(16, 1, "現價（2026-10-23 收盤，Euronext Amsterdam）"); V["B16"] = "=px"; V["B16"].number_format = "#,##0.00"
for rr, y, lab in ((17, 2026, "現價 ÷ FY2026 EPS（錨點）"), (18, 2027, "現價 ÷ FY2027 EPS（模型）"), (19, 2028, "現價 ÷ FY2028 EPS（模型）"), (20, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(rr, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{rr}"] = f"=px/PL!{FYC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(21, 1, "前瞻 P/E（stockanalysis）"); V["B21"] = "=sa_fpe"; V["B21"].number_format = "0.0"
V.cell(22, 1, "現價 ÷ 共識 FY2027 EPS"); V["B22"] = "=px/cons_eps27"; V["B22"].number_format = "0.0"
V["E17"] = "FY2026 三情境相同（錨點）"; V["E21"] = "近 12 月 P/E 40.24（同站）"

V.cell(24, 1, "2. 10 年 P/E（年底）").font = BOLD
hdr(V, 25, ["年底", "ADR 收盤價 US$（未調整）", "基本 EPS €", "ECB 年底 US$/€", "EPS 換美元", "P/E 自算", "companiesmarketcap P/E", "自算 − cmc", "標記", "來源／說明"])
PXA = {2016: 112.20, 2017: 173.82, 2018: 155.62, 2019: 295.94, 2020: 487.72, 2021: 796.14, 2022: 546.40, 2023: 756.92, 2024: 693.08, 2025: 1069.86}
EPSA = {2016: 3.46, 2017: 4.93, 2018: 6.10, 2019: 6.16, 2020: 8.49, 2021: 14.36, 2022: 14.14, 2023: 19.91, 2024: 19.25, 2025: 24.73}
FXE = {2016: 1.0541, 2017: 1.1993, 2018: 1.1450, 2019: 1.1234, 2020: 1.2271, 2021: 1.1326, 2022: 1.0666, 2023: 1.1050, 2024: 1.0389, 2025: 1.1750}
CMC = {2016: 27.3, 2017: 28.4, 2018: 20.4, 2019: 41.0, 2020: 46.9, 2021: 46.1, 2022: 36.2, 2023: 34.6, 2024: 33.4, 2025: 37.5}
EPSSRC = {2016: "Q4 2017 新聞稿（6-K）：2016 基本 3.46", 2017: "同上：2017 基本 4.93", 2018: "Q4 2019 新聞稿（6-K）：2018 基本 6.10", 2019: "同上：2019 基本 6.16",
          2020: "Q4 2020 新聞稿（6-K）：基本 8.49", 2021: "Q4 2022 新聞稿（6-K）：2021 基本 14.36", 2022: "同上：2022 基本 14.14", 2023: "Q4 2024 新聞稿（6-K）：2023 基本 19.91",
          2024: "同上：2024 基本 19.25", 2025: "Q4 2025 簡報（6-K）：基本 24.73"}
for i, y in enumerate(range(2016, 2026), 26):
    V.cell(i, 1, f"{y}-12")
    for col, v in (("B", PXA[y]), ("C", EPSA[y]), ("D", FXE[y]), ("G", CMC[y])):
        V[f"{col}{i}"] = v; V[f"{col}{i}"].font = BLUE
    V[f"E{i}"] = f"=C{i}*D{i}"; V[f"F{i}"] = f"=B{i}/E{i}"; V[f"H{i}"] = f"=F{i}-G{i}"
    for c, fm in (("B", "#,##0.00"), ("C", "0.00"), ("D", "0.0000"), ("E", "0.00"), ("F", "0.0"), ("G", "0.0"), ("H", "0.0")): V[f"{c}{i}"].number_format = fm
    V.cell(i, 9, "Verified／Interested-party"); V.cell(i, 10, f"價格：digrin（未調整 Real price）；EPS：{EPSSRC[y]}；匯率：ECB 參考匯率")
V.cell(36, 1, "來源：ADR 年底收盤價＝digrin.com（Real price；Macrotrends 的調整後價格作數量級對照）；EPS＝ASML 各年第四季新聞稿（SEC 6-K），2021–2025 另以 stockanalysis 稀釋 EPS（14.34、14.13、19.89、19.24、24.71）交叉；匯率＝ECB 各年最後營業日參考匯率。cmc（companiesmarketcap）為美元計價 TTM 口徑，作第二來源對照（差距見 H 欄）。Euronext 年底收盤價：未找到（改用 ADR，口徑相同）。")
V.cell(38, 1, "3. 分位").font = BOLD
hdr(V, 39, ["統計", "P/E 自算", "cmc P/E（參考）", "用途"])
for i, (nm, p, use) in enumerate((("25 分位", 0.25, "見頂年"), ("中位數", 0.5, "中段"), ("75 分位", 0.75, "谷底")), 40):
    V.cell(i, 1, nm); V.cell(i, 4, use)
    V.cell(i, 2, f"=PERCENTILE(F26:F35,{p})"); V.cell(i, 3, f"=PERCENTILE(G26:G35,{p})")
    for c in "BC": V[f"{c}{i}"].number_format = "0.0"
V.cell(43, 1, "年數"); V["B43"] = "=COUNT(F26:F35)"; V["C43"] = "=COUNT(G26:G35)"
defname("pe_peak", "Val!$B$40"); defname("pe_mid", "Val!$B$41"); defname("pe_trough", "Val!$B$42")
V.cell(44, 1, "讀法：2016–2018 年 P/E 22–31x（EUV 量產前）、2019–2021 年 43–49x（EUV 量產與疫情後設備熱潮）、2022–2025 年 34–37x。分位對樣本敏感；兩來源的分位差距見 C 欄。")

V.cell(46, 1, "4. 循環位置（設備股共用規則；見頂年度讀設備市場模型，不在本模型另訂）").font = BOLD
V.cell(47, 1, CP.rule_text(PEAK) + "。ASML 會計年度＝日曆年，FY(t) 全部落在 t 年")
hdr(V, 48, ["情境", "設備市場見頂年度", "FY2027 對照日曆年", "判定", "倍數", "規則"])
for i, s in enumerate(SC, 49):
    V.cell(i, 1, s); V.cell(i, 2, f"=pk_{s}"); V.cell(i, 3, 2027).font = BLUE
    V.cell(i, 4, "=" + CP.excel_rule(f"C{i}", f"pk_{s}"))
    V.cell(i, 5, f'=IF(D{i}="見頂",pe_peak,IF(D{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 6, "tools/cycle_position.py")
V.cell(53, 1, "5. 隱含股價（歐元）：規則值＝FY2027 EPS × 依循環位置的倍數；其餘倍數僅供參考").font = BOLD
hdr(V, 54, ["情境", "EPS 年度", "EPS", "倍數類型", "倍數", "隱含股價", "相對現價", "備註"])
vr = 55; VAL = {}
for s in SC:
    for y, typ, nm in ((2027, f"E{49 + SC.index(s)}", "rule"), (2027, "pe_peak", "參考：25 分位"), (2027, "pe_mid", "參考：中位數"), (2027, "pe_trough", "參考：75 分位"),
                       (2028, "pe_peak", "參考：25 分位"), (2028, "pe_mid", "參考：中位數"), (2028, "pe_trough", "參考：75 分位")):
        V.cell(vr, 1, s); V.cell(vr, 2, f"FY{y}"); V.cell(vr, 3, f"=PL!{FYC[y]}{SUMR[(s,'eps')]}").number_format = "#,##0.00"
        V.cell(vr, 4, nm if nm != "rule" else f'="規則值："&D{49 + SC.index(s)}'); V.cell(vr, 5, f"={typ}").number_format = "0.0"
        V.cell(vr, 6, f"=C{vr}*E{vr}").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
        V.cell(vr, 8, "未折現；約 FY2027 年底的價格" if nm == "rule" else "")
        if nm == "rule":
            for c in range(1, 9): V.cell(vr, c).fill = KEY
            VAL[s] = vr
        vr += 1
V.cell(vr + 1, 1, "本版未做先驗：無 Andy 合理 P/E 對照。")
# 0. 目標價與上檔空間（Andy 2026-10-25 要求；列標籤依工作單逐字）
V.cell(4, 1, "0. 目標價與上檔空間").font = BOLD
hdr(V, 5, ["項目", "數值", "", "", "說明"])
V.cell(6, 1, "目標價（12 個月，情境加權）"); V["B6"] = f"=tp_wA*F{VAL['A']}+tp_wB*F{VAL['B']}+tp_wC*F{VAL['C']}"; V["B6"].number_format = "#,##0"
V["E6"] = "權重 × 各情境隱含價（第 5 節規則值）"
V.cell(7, 1, "上檔空間（目標價 ÷ 現價 − 1）"); V["B7"] = "=B6/px-1"; V["B7"].number_format = "0.0%"; V["E7"] = "現價＝In 頁 px（2026-10-23 收盤）"
for rr, s, lab in ((8, "A", "情境 A 隱含價（樂觀）"), (9, "B", "情境 B 隱含價（基準）"), (10, "C", "情境 C 隱含價（下檔）")):
    V.cell(rr, 1, lab); V[f"B{rr}"] = f"=F{VAL[s]}"; V[f"B{rr}"].number_format = "#,##0"
    V[f"E{rr}"] = f'="FY2027 EPS × "&D{49 + SC.index(s)}&" 倍數（第 5 節）"'
V.cell(11, 1, "情境權重 A／B／C"); V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0%"
V["E11"] = '="合計 "&TEXT(B11+C11+D11,"0%")&"；In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"'
V.cell(12, 1, "目標價方法")
V["B12"] = ("P/E：各情境 FY2027 EPS × 十年（2016–2025 年底）P/E 依循環位置取分位（設備股共用規則：A、B 中段＝中位數，C 見頂＝25 分位）。"
            "十年倍數含 2019–2021 年 EUV 量產與疫情後設備熱潮的高倍數（43–49x），偏誤方向為高估；隱含價約為 FY2027 年底的價格、未折現，當作 12 個月目標價時偏高")
V.column_dimensions["A"].width = 46
for c in "BCDEFGHI": V.column_dimensions[c].width = 15
V.column_dimensions["J"].width = 90

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A）：對 FY2027、FY2028 營收與 EPS 的影響。每個案例是一個完整的推導區塊（本頁下方，與 PL 同一引擎）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收 €B", "FY2027 EPS €", "相對基準", "FY2028 營收 €B", "FY2028 EPS €", "相對基準"])
cases = [("微影強度與價格 0%／年", "g_int＝0", dict(gi="g_int_lo")), ("微影強度與價格 +6%／年", "g_int＝6%", dict(gi="g_int_hi")),
         ("中國占 FY2026 系統 17%", "非中國基期較大", dict(cn="cn_sh26_lo")), ("中國占 FY2026 系統 23%", "非中國基期較小", dict(cn="cn_sh26_hi")),
         ("中國年成長 −10%", "情境 A 基準 0%", dict(gcn="g_cn_A_lo")), ("中國年成長 +10%", "", dict(gcn="g_cn_A_hi")),
         ("IBM 年成長 +8%", "基準 +12%", dict(gibm="g_ibm_A_lo")), ("IBM 年成長 +16%", "", dict(gibm="g_ibm_A_hi")),
         ("毛利率斜率 ×0.5", "營收規模帶動毛利率的效果減半", dict(gmb="gm_b_lo")), ("毛利率斜率 ×1.5", "", dict(gmb="gm_b_hi")),
         ("營業費用全固定（phi＝0）", "", dict(phi="phi_lo")), ("營業費用變動比例 40%", "營業費用隨營收成長較多", dict(phi="phi_hi")),
         ("2027 產能 +20%", "Low-NA「至少 80 台」的下緣口徑", dict(c27="cap_27_lo")), ("2028 產能 0%", "", dict(c28="cap_28_lo")),
         ("淨利÷營業利益 0.80", "", dict(nioi="nioi_lo")), ("淨利÷營業利益 0.87", "", dict(nioi="nioi_hi")),
         ("買回 0%／年", "", dict(bb="bb_lo")), ("買回 2%／年", "", dict(bb="bb_hi"))]
def put_case(i, name, setting, R, pre=""):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    SE[f"C{i}"] = f"={pre}{FYC[2027]}{R['rev']}"; SE[f"D{i}"] = f"={pre}{FYC[2027]}{R['eps']}"; SE[f"F{i}"] = f"={pre}{FYC[2028]}{R['rev']}"; SE[f"G{i}"] = f"={pre}{FYC[2028]}{R['eps']}"
    SE[f"E{i}"] = f"=D{i}/D$4-1"; SE[f"H{i}"] = f"=G{i}/G$4-1"
    for c in "CF": SE[f"{c}{i}"].number_format = "#,##0.0"
    for c in "DG": SE[f"{c}{i}"].number_format = "#,##0.00"
    for c in "EH": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
SB = {}
for n, (nm, setting, ov) in enumerate(cases):
    SB[nm] = engine(SE, 50 + n * 40, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm])
nn = 5 + len(cases)
put_case(nn, "情境 B", "設備支出 2028 持平；中國 −10%；IBM +8%", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "設備支出 2028 −20%；中國 −25%；IBM +3%", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 三個設備情境相同（設備市場模型 2027 年三情境相同），B、C 在 FY2027 的差異只來自中國與 IBM；設備情境的差異從 FY2028 才出現。產能上限在基準下未生效（Recon），所以產能案例對 EPS 影響為 0 時，代表需求低於產能。")
SE.column_dimensions["A"].width = 34; SE.column_dimensions["B"].width = 36
for c in "CDEFGH": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、共識、公司 2030 目標、設備市場與產能。只作事後對照，不為對上數字而調整參數。差距 > 10% 的原因見 README"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False):
    global rr_
    RC.cell(rr_, 1, label); RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        L = chr(64 + cv)
        RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y=2026: (lambda s: f"=PL!{FYC[y]}{SUMR[(s, key)]}")
eg = lambda key, y: (lambda s: f"=PL!{FYC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. FY2026 錨點 vs 公司全年指引（錨點用第四季指引，故差距小；非獨立檢驗）").font = BOLD; rr_ += 1
rec("營收（€B）", "=g_fy_rev/1000", pl("rev"), "Q3 2026 新聞稿：€44–45B 中點")
rec("毛利率（差距以百分點計）", "=g_fy_gm", pl("gm"), "同上：約 56%", fmt="0.0%", pt=True)
rec("IBM 成長（差距以百分點計）", "=g_ibm26", lambda s: f"=PL!D{PB[s]['ibm']}/PL!C{PB[s]['ibm']}-1", "Q3 2026 法說：約 +50%", fmt="0.0%", pt=True)
rec("系統營收（€B）vs 指引成長推得（邏輯 +10%、記憶體 +75%）", "=(a25_log*(1+g_log)+a25_mem*(1+g_mem))/1000", pl("sys"), "Q3 2026 法說", "獨立檢驗：錨點系統營收（由第四季總營收與 IBM 指引推得）vs 邏輯／記憶體成長指引")
RC.cell(rr_, 1, "B. vs 共識（stockanalysis 2026-10-17；eulerpool）").font = BOLD; rr_ += 1
rec("FY2026 營收（€B）", "=cons_rev26/1000", pl("rev"), "stockanalysis 44.46B")
rec("FY2026 EPS（€）", "=cons_eps26", pl("eps"), "stockanalysis 41.21（non-GAAP；eulerpool 41.24）", "模型為報告口徑，含 Mistral 評價益約 €1.3／股", fmt="0.00")
rec("FY2027 營收（€B）", "=cons_rev27/1000", pl("rev", 2027), "stockanalysis 57.21B")
rec("FY2027 EPS（€）", "=cons_eps27", pl("eps", 2027), "stockanalysis 55.93（eulerpool 隱含 56.7）", fmt="0.00")
rec("FY2028 EPS（€；eulerpool P/E 18.80 隱含）", "=px/eul_pe28", pl("eps", 2028), "eulerpool（單一來源）", fmt="0.00")
RC.cell(rr_, 1, "C. FY2029 vs 公司 2030 目標（2024 投資人日；目標年較模型末年晚一年）").font = BOLD; rr_ += 1
rec("營收（€B）vs 目標下緣 44", "=t30_rev_lo/1000", pl("rev", 2029), "公司：€44–60B")
rec("營收（€B）vs 目標上緣 60", "=t30_rev_hi/1000", pl("rev", 2029), "同上")
rec("毛利率 vs 目標中點 58%（百分點）", "=(t30_gm_lo+t30_gm_hi)/2", pl("gm", 2029), "公司：56–60%", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "D. 微影占晶圓前段設備（ASML 系統營收換美元 ÷ 設備市場模型；FY2025 對照 SEMI 實績）").font = BOLD; rr_ += 1
rec("FY2025：ASML 占 SEMI 晶圓前段設備實績（百分點）", "=a25_sys/1000*fx_25/semi_wfe_25", eg("lsh", 2025), "SEMI 2025 實績（設備市場模型）", "設備市場 2025 以 SEMI 實績校準，差距必為 0（恆等式）；比例本身約 23.7%", fmt="0.0%", pt=True)
for y in (2026, 2027, 2028, 2029):
    rec(f"FY{y}：ASML 占晶圓前段設備（各情境）", f"=PL!{FYC[2025]}{PB['A']['lsh']}", eg("lsh", y), "對照值＝FY2025 比例", "比例上升＝微影強度與價格（g_int）＋記憶體比重上升；FY2026 起匯率 Assumed", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "E. 產能與 EUV 台數").font = BOLD; rr_ += 1
rec("FY2027 隱含 Low-NA EUV 台數 vs 公司 2027 產能「至少 80 台」", "=euv_cap27", eg("lna", 2027), "Q1 2026 法說", "由 EUV 占比、High-NA 台數與單價（皆 Assumed）推得；超過產能代表 EUV 占比或單價假設偏高", fmt="0")
for y in (2027, 2028, 2029):
    rec(f"FY{y}：系統需求 ÷ 產能上限", "=1", (lambda yy: (lambda s: f"=PL!{FYC[yy]}{PB[s]['dem']}/PL!{FYC[yy]}{PB[s]['cap']}"))(y), "產能上限＝FY2026 × 產能成長 × 強度係數", ">100%＝產能限制生效（系統營收取上限）", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "F. 中國占比路徑").font = BOLD; rr_ += 1
rec("FY2025 中國占總營收 33% vs 模型 FY2029 中國占系統營收（百分點）", "=cn_tot25", pl("cnsh", 2029), "Caixin（公司揭露）", "口徑不同（總營收 vs 系統），只看方向", fmt="0.0%", pt=True)
RC.column_dimensions["A"].width = 64
for c in "BCDEFGH": RC.column_dimensions[c].width = 11
RC.column_dimensions["I"].width = 38; RC.column_dimensions["J"].width = 80
RECON_LAST = rr_ - 1

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（SEC 6-K 為公司原文；asml.com 的 PDF 與第三方頁面經搜尋引擎摘錄或擷取讀取）"; SO["A1"].font = BOLD
srcs = [("公司：Q3 2026 新聞稿（SEC 6-K，2026-10-14）", "https://www.sec.gov/Archives/edgar/data/0000937966/000162828026065813/pressreleasefinancialresul.htm"),
 ("公司：Q3 2026 財務報表（季度摘要 Q3 2025–Q3 2026、9M）", "https://www.sec.gov/Archives/edgar/data/0000937966/000162828026065813/financial_statementsxusx.htm"),
 ("公司：Q3 2026 簡報（第四季與全年 2026 指引、IBM、稅率）", "https://www.sec.gov/Archives/edgar/data/0000937966/000162828026065813/a2026_10x14xpresentation.htm"),
 ("Q3 2026 法說逐字稿（Benzinga；EUV €4.1B、邏輯 54%、中國約 20%、Mistral 評價益）", "https://www.benzinga.com/news/26/10/62356129"),
 ("stockhub.kr：Q3 2026 摘要（營業利益 €5.2B、稀釋 EPS 12.93）", "https://www.stockhub.kr/en/insights/asml-cy2026-q3-earnings-20261014"),
 ("aiweekly：Q3 2026 中國占系統出貨 23%（Q2 14%）", "https://aiweekly.co/alerts/asml-q3-china-equipment-sales-soar-109-qoq-to-23-of-system-shipments-as-memory"),
 ("公司：Q2 2026 法說逐字稿（2026-07-15）", "https://ourbrand.asml.com/asset/1fd3908a-0381-47b5-9b69-a3094f656651/2026_07_15-ASML-Transcript-investor-call-Q2-2026.pdf"),
 ("公司：Q1 2026 法說逐字稿（2026-04-15；2027 Low-NA 至少 80 台）", "https://ourbrand.asml.com/asset/8e1f7393-33dd-4737-a436-cfe1b68cc577/2026_04_15-ASML-Transcript-investor-call-Q1-2026.pdf"),
 ("公司：Q4 2025 簡報（SEC 6-K，2026-01-28；FY2025 EUV、DUV、IBM、買回 €12B）", "https://www.sec.gov/Archives/edgar/data/937966/000162828026003701/a2026_01x28xpresentation.htm"),
 ("公司：Q4 2025 法說逐字稿（邏輯 €16.1B、記憶體 €8.4B）", "https://ourbrand.asml.com/m/285dcba22c2bf203/original/2026_01_28-ASML-Transcript-investor-call-Q4-2025.pdf"),
 ("Caixin：ASML 2025 中國占 33%、2026 約 20%（2026-01-29）", "https://www.caixinglobal.com/2026-01-29/asml-expects-china-revenue-drop-following-backlog-fueled-surge-102409285.html"),
 ("公司：Q4 2024 新聞稿（6-K；2023、2024 EPS）", "https://www.sec.gov/Archives/edgar/data/937966/000093796625000003/pressreleasequarterlyresul.htm"),
 ("公司：Q4 2022 新聞稿（6-K；2021、2022 EPS）", "https://www.sec.gov/Archives/edgar/data/937966/000093796623000003/pressreleasequarterlyresul.htm"),
 ("公司：Q4 2020 新聞稿（6-K；2020 EPS）", "https://www.sec.gov/Archives/edgar/data/937966/000093796621000003/pressreleasequarterlyresul.htm"),
 ("公司：Q4 2019 新聞稿（6-K；2018、2019 EPS）", "https://www.sec.gov/Archives/edgar/data/937966/000093796620000003/pressreleasequarterlyresul.htm"),
 ("公司：Q4 2017 新聞稿（6-K；2016、2017 EPS）", "https://www.sec.gov/Archives/edgar/data/0000937966/000093796618000003/pressreleasequarterlyresul.htm"),
 ("stockanalysis：ASML 財務（FY2021–2025，稀釋 EPS 交叉）", "https://stockanalysis.com/stocks/asml/financials/"),
 ("stockanalysis：AMS:ASML 現價、前瞻 P/E（2026-10-23）", "https://stockanalysis.com/quote/ams/ASML/"),
 ("stockanalysis：共識（2026-10-17 更新）", "https://stockanalysis.com/stocks/asml/forecast/"),
 ("eulerpool：ASML.AS 現價、共識 P/E", "https://eulerpool.com/stock/NL0010273215"),
 ("digrin：ASML ADR 年底價格（未調整）", "https://www.digrin.com/stocks/detail/ASML/price"),
 ("Macrotrends：ASML ADR 價格歷史（調整後，數量級對照）", "https://macrotrends.net/stocks/charts/ASML/asml-holding/stock-price-history"),
 ("companiesmarketcap：ASML 年底 P/E", "https://companiesmarketcap.com/asml/pe-ratio/"),
 ("ECB：歐元參考匯率（各年最後營業日，例：2025-12-31）", "https://www.ecb.europa.eu/stats/exchange/eurofxref/shared/pdf/2025/12/20251231.pdf"),
 ("芬蘭央行：2025 年平均匯率（ECB）", "https://www.suomenpankki.fi/en/statistics/interest-rates-and-exchange-rates/exchange-rates/average-annual-exchange-rates-for-the-year-2025/"),
 ("Yole／Tom's Hardware：High-NA 約 US$380M", "https://www.yolegroup.com/industry-news/asmls-high-na-chipmaking-tool-will-cost-380-million-the-company-already-has-orders-for-10-to-20-machines-and-is-ramping-up-production/"),
 ("市場層：markets/equipment 成品（本 repo；晶圓前段設備三情境、見頂年度）", rel(EQ_F)),
 ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 70; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("ASML（ASML.AS／ASML ADR）收支與估值模型 v0.2（AI 半導體 Project 步驟 5：設備與材料）　2026-10-25", True),
 ("模型命題：ASML 是 EUV 唯一供應者，營收由晶圓廠的微影設備支出與裝機基礎服務決定。2026 年 AI 帶動記憶體與先進邏輯擴產，營收約 +36%、毛利率升到 56%；股價約 25 倍前瞻本益比，問題是設備支出在 2027 年後放緩、中國營收占比續降時，產能 +30% 與漲價能否撐住盈餘成長。", False), ("", False),
 ("會計年度＝日曆年（FY2026＝2026-01～12）；金額以歐元為主（Q 頁 €M，PL 頁 €B）。FY2026 錨點＝前三季實績（含 2026-10-14 公布的第三季）＋第四季公司指引中點，三情境相同；FY2027–FY2029 由驅動推導。第四季財報預定 2027-01 下旬公布。", False), ("", False),
 ("驅動因子 → 推導量", True),
 ("系統營收（非中國）＝FY2026 錨點 × 晶圓前段設備指數（markets/equipment：邏輯＝先進邏輯＋其他代工邏輯，記憶體＝DRAM 含 HBM＋NAND；FY2026＝1）× 微影強度與價格係數（Assumed +3%／年）；邏輯與記憶體分開，權重＝FY2026 指引成長推得的邏輯占比", False),
 ("中國系統營收：FY2026＝錨點 × 20%（公司：2026 年約 20%），之後 × (1＋情境別年成長)；系統營收＝MIN(需求, 產能上限)，產能＝FY2026 ×（2027 +30%、2028 +15%、2029 +10%）× 強度係數", False),
 ("IBM：前一年 × (1＋情境別成長)。毛利率＝FY2026 錨點毛利率＋斜率 × ln(營收比)（FY2025 與 FY2026 兩點推導）；營業費用＝FY2026 ×［固定＋變動 × 營收比］（兩點推導）；不寫死利益率", False),
 ("情境 A／B／C：設備市場模型的三情境（A 續增；B 2028 持平、2029 −5%；C 2028 −20%）＋中國年成長（0%／−10%／−25%）＋IBM 年成長（+12%／+8%／+3%）", False),
 ("Val：P/E 規則——10 年（2016–2025 年底）P/E 分位，見頂年 25 分位、中段中位數、谷底 75 分位；循環位置依設備股共用規則（各情境見頂年度讀設備市場模型：" + "、".join(f"{s_} {PEAK[s_]}" for s_ in SC) + " → FY2027：" + CP.fy2027_text(PEAK) + "）。本版未做先驗，無 Andy 合理 P/E 對照。Recon：公司指引、共識、2030 目標、微影占設備支出、產能。Sens：強度與價格、中國、IBM、毛利率斜率、營業費用等", False), ("", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。實績取自 SEC 6-K；法說內容取自公司逐字稿與第三方轉述。", False),
 ("已知限制：(1) 設備市場的其他代工與記憶體廠含中國廠，本模型另以中國系統營收列處理，指數的中國成分重疊；(2) 微影強度與價格（g_int）無直接實證，只有 FY2025 一年可比；(3) 毛利率、營業費用都只有兩點，毛利率對規模的斜率可能高估（FY2026 的 +2.9pt 含漲價與產品組合）；(4) 指數只連結設備支出的成長，匯率變動不進營收；(5) 產能上限以系統總額近似，未分 EUV／DUV；(6) Euronext 年底收盤價未找到，10 年 P/E 用 ADR 價格與 ECB 匯率換算；(7) 2028–29 依設備市場情境係數（Assumed）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

CFG_V2 = dict(cols=FYC, scn={"A": "順利", "B": "平平", "C": "不如預期"}, peers=[('應用材料（Applied Materials）', 'AMAT', 21.39, '2026-10-23', 'stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）'),('科林研發（Lam Research）', 'LRCX', 22.39, '2026-10-23', 'stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）'),('科磊（KLA）', 'KLAC', 26.65, '2026-10-23', 'stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）'),('東京威力科創（Tokyo Electron）', '8035.T', 27.09, '2026-10-21', 'stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）'),('ASM International', 'ASM.AS', 25.77, '2026-10-23', 'stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）')], fwd={2027: "2/12", 2028: "10/12"}, fwd_note='約 2027-10 時的未來 12 月 ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）', a3='目標價（DCF 與 Comparable 加權）：見本頁第 {row} 列起「0. 目標價」（估值 v2 新增，放在既有內容之後以免位移既有列）；原「0. 目標價與上檔空間」改名為「舊法：規則值情境加權（參考）」',
    cur='€',
    unit_rev='€B',
    unit_ni='€B',
    fmt='#,##0.00',
    fmt_rev='#,##0.0',
    band_fmt='#,##0',
    roe=('FY2023–FY2025', [0.7042, 0.4743, 0.5046], 'stockanalysis 比率頁（2026-10-25 讀取）：FY2023 70.42%、FY2024 47.43%、FY2025 50.46%', '單一來源；近 12 月 65.64%（同站）'),
    nim=('FY2016–FY2025', [(6.88, 1.56), (8.96, 2.07), (10.94, 2.59), (11.82, 2.59), (13.98, 3.55), (18.61, 5.88), (21.17, 5.62), (27.56, 7.84), (28.26, 7.57), (32.67, 9.61)], 'vcpscanner.com 年度損益（十億；網站標示美元，數值與公司歐元報表相同量級；比率不受幣別影響）', '第二來源：FY2025 與 Q 頁實績核對；其餘年度單一來源'),
    cons_tp=(2108, 'stockanalysis（AMS:ASML）：41 位分析師平均 €2,108（中位數 2,100；區間 1,288–2,600），資料更新 2026-10-21', '第二來源：未找到'),
    cons_feps=('=cons_eps27', '共識前瞻 EPS＝共識 FY2027 EPS（In 頁 cons_eps27，stockanalysis 55.93）；本模型前瞻 EPS 以 FY2028 為主，約晚一年，EPS 因素偏大', 'Derived'),
    peer_note='前段設備同業：應用材料、科林研發、科磊、東京威力科創、ASM International',
    dcf_note='',
    pe10='F26:F35',
    pe10_note='現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（自算）的分位；只作檢查，不進目標價',
    px_note='現價＝In 頁 px（Euronext 2026-10-23 收盤）',
    method='DCF：股東現金流＝EPS ×（1 − g／ROE）（模型無折舊與資本支出項目），r 10%、g 4%、ROE 為近三年平均，FY2027–32 明確期（FY2030 起收斂法）＋永續。Comparable：應用材料、科林研發、科磊、東京威力科創、ASM International 前瞻 P/E 中位數 × 12 個月後前瞻 EPS（同業多為美元計價，倍數不受幣別影響）。偏誤：ASML 的 EUV 獨占通常享有溢價，用同業中位數偏低估；十年平均淨利率約兩成六，情境 C 收斂偏保守。',
    readme=['v0.2（2026-10-25）估值 v2：Val 頁新增「0. 目標價」（放在既有內容之後，不位移既有列）＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，情境權重 A 0.3／B 0.5／C 0.2）＋0.5 × Comparable 目標價（同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解、歷史分位檢查。原「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」，不再叫目標價。', '預估期間延伸到 FY2032：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（PL 頁底部延伸區塊，灰底）。情境對應：A＝順利、B＝平平、C＝不如預期（原為顆數情境）。本版未做先驗，無 Andy 合理 P/E 對照。'])

# ---------------- 估值 v2（2026-10-25，Andy 11:18／11:20／11:28／11:53）：PL 延伸到 FY2032（收斂法）＋ Val「0. 目標價」（DCF 與 Comparable 加權） ----------------
# 為不位移既有儲存格（FY2026–FY2029 既有 EPS 逐格不變），新內容放在各頁既有內容之後；HTML 估值章第一個區塊讀「0.3 結論」
def valuation_v2(CFG):
    cols = CFG["cols"]; MF = CFG["fmt"]; YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
    FADEFILL = PatternFill("solid", fgColor="E7E6E6")
    # --- In：目標價與收斂法參數（黃底；Assumed，Andy 可改） ---
    r = IN.max_row + 2
    IN.cell(r, 1, "J. 目標價（DCF 與 Comparable 加權）與 FY2030–32 收斂法參數（v2 新增；黃底＝Assumed，Andy 可改）").font = BOLD; r += 1
    def inrow(key, item, unit, base, tag, s1, s2, fmt=None, yellow=False):
        nonlocal r
        IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
        c = IN.cell(r, 5, base)
        if not (isinstance(base, str) and base.startswith("=")): c.font = BLUE
        if fmt: c.number_format = fmt
        if yellow: c.fill = KEY
        IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
        wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
        r += 1
    inrow("tp_r", "折現率（股東權益要求報酬；美、台、日、韓同一值）", "%", 0.10, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%）", CFG.get("r_note", "—"), "0.0%", True)
    inrow("tp_g", "長期成長率（永續成長；收斂法 FY2032 營收年增的基準）", "%", 0.04, "Assumed", "Andy 可改（工作單預設 4%）", "情境 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt（PL 頁延伸區塊）", "0.0%", True)
    inrow("tp_roe", f"ROE（近三年平均，{CFG['roe'][0]}）", "%", "=AVERAGE(" + ",".join(str(x) for x in CFG["roe"][1]) + ")", "Derived", CFG["roe"][2], CFG["roe"][3], "0.0%")
    inrow("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", 0.5, "Assumed", "Andy 可改（2026-10-25 11:53：DCF／Comparable 各半）", "—", "0%", True)
    if CFG["nim"]:
        inrow("nim10", f"十年平均淨利率（{CFG['nim'][0]}，各年淨利 ÷ 營收的平均）", "%", "=AVERAGE(" + ",".join(f"{n}/{d}" for d, n in CFG["nim"][1]) + ")", "Derived", CFG["nim"][2], CFG["nim"][3], "0.0%")
    else:
        IN.cell(r, 1, "nim10"); IN.cell(r, 2, "十年平均淨利率"); IN.cell(r, 3, "%"); IN.cell(r, 5, "未找到"); IN.cell(r, 7, "Assumed")
        IN.cell(r, 8, CFG["nim_missing"]); IN.cell(r, 9, "依工作單：找不到十年平均時用最後一年（FY2029）淨利率 × 0.8；PL 延伸區塊逐情境計算"); r += 1
    inrow("cons_tp", "市場共識目標價（賣方平均）", CFG["cur"], CFG["cons_tp"][0], "Interested-party", CFG["cons_tp"][1], CFG["cons_tp"][2], MF)
    inrow("cons_feps", "共識前瞻 EPS（差距拆解用）", CFG["cur"], CFG["cons_feps"][0], "Interested-party", CFG["cons_feps"][1], CFG["cons_feps"][2], MF)

    # --- PL：延伸到 FY2032（收斂法） ---
    p0 = PL.max_row + 3
    PL.cell(p0, 1, "延伸到 FY2032（v2 新增）：FY2025–FY2029 引用本頁頂端摘要；灰底 FY2030–FY2032 為收斂法，非驅動因子推導").font = BOLD
    PL.cell(p0 + 1, 1, "收斂法：營收年增由 FY2029 年增線性收斂到 FY2032 的長期成長（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt）；淨利率由 FY2029 線性收斂（A＝FY2029、B＝FY2029 與十年平均的中點、C＝十年平均" + ("" if CFG["nim"] else "；十年平均未找到，以 FY2029 × 0.8 代替（Assumed）") + "）；股數維持 FY2029（EPS 與淨利同比例）")
    ext_cols = {y: chr(ord("C") + y - 2025) for y in range(2025, 2033)}
    hdr(PL, p0 + 2, ["情境（延伸）", "單位"] + [str(PL[f"{cols[y]}3"].value) for y in range(2025, 2030)] + ["FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    EXT = {}; pr = p0 + 3
    gadd = {"A": "+0.02", "B": "", "C": "-0.02"}
    for s in SC:
        PL.cell(pr, 1, f"情境 {s}（{CFG['scn'][s]}）").font = BOLD; pr += 1
        R = {k: pr + i for i, k in enumerate(("rev", "g", "nim", "ni", "eps"))}
        for k, lab in (("rev", "營收"), ("g", "  營收年增"), ("nim", "  淨利率（淨利 ÷ 營收）"), ("ni", "淨利（收斂法延伸）"), ("eps", "EPS（延伸）")):
            PL.cell(R[k], 1, lab); PL.cell(R[k], 2, {"rev": CFG["unit_rev"], "ni": CFG["unit_ni"], "eps": CFG["cur"]}.get(k, "%"))
        for y in range(2025, 2033):
            c = ext_cols[y]; pc = ext_cols.get(y - 1)
            if y <= 2029:
                PL[f"{c}{R['rev']}"] = f"={cols[y]}{SUMR[(s, 'rev')]}"
                PL[f"{c}{R['ni']}"] = f"={cols[y]}{SUMR[(s, 'ni')]}"
                PL[f"{c}{R['eps']}"] = f"={cols[y]}{SUMR[(s, 'eps')]}"
                PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
                if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            else:
                k = y - 2029
                mT = {"A": f"$G{R['nim']}", "B": f"($G{R['nim']}+{'nim10' if CFG['nim'] else '0.8*$G' + str(R['nim'])})/2", "C": "nim10" if CFG["nim"] else f"0.8*$G{R['nim']}"}[s]
                PL[f"{c}{R['g']}"] = f"=$G{R['g']}+((tp_g{gadd[s]})-$G{R['g']})*{k}/3"
                PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
                PL[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({mT}-$G{R['nim']})*{k}/3"
                PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
                PL[f"{c}{R['eps']}"] = f"=$G{R['eps']}*{c}{R['ni']}/$G{R['ni']}"
                for kk in R: PL[f"{c}{R[kk]}"].fill = FADEFILL
            for kk, fm in (("rev", CFG["fmt_rev"]), ("g", "0.0%"), ("nim", "0.0%"), ("ni", CFG["fmt_rev"]), ("eps", MF)):
                if PL[f"{c}{R[kk]}"].value is not None: PL[f"{c}{R[kk]}"].number_format = fm
        PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導")
        EXT[s] = R; pr += 6
    for c in "IJ": PL.column_dimensions[c].width = 14
    PL.column_dimensions["K"].width = 40

    # --- Val：舊法改名（不再叫目標價） ---
    REN = (("0. 目標價與上檔空間", "舊法：規則值情境加權（參考）"), ("目標價（12 個月，情境加權）", "規則值情境加權價（12 個月；參考）"),
           ("上檔空間（目標價 ÷ 現價 − 1）", "相對現價（規則值情境加權價 ÷ 現價 − 1）"), ("目標價方法", "規則值方法"))
    for row in V.iter_rows(min_col=1, max_col=1):
        v = row[0].value
        if isinstance(v, str):
            for a, b in REN:
                if v.startswith(a): row[0].value = b + v[len(a):]; break
    # --- Val：0. 目標價（列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期） ---
    t0 = V.max_row + 3
    T = {}
    def tset(rr, key, label, fs, fmt, note=""):
        V.cell(rr, 1, label); T[key] = rr
        for c, f in zip("BCD", fs):
            if f is None: continue
            V[f"{c}{rr}"] = f; V[f"{c}{rr}"].number_format = fmt
        if note: V[f"E{rr}"] = note
    V.cell(t0, 1, "0. 目標價").font = BOLD; T["top"] = t0
    V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD; T["h01"] = t0 + 1
    hdr(V, t0 + 2, ["項目", f"A（{CFG['scn']['A']}）", f"B（{CFG['scn']['B']}）", f"C（{CFG['scn']['C']}）", "說明"]); T["hd1"] = t0 + 2
    V.cell(t0 + 3, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{t0 + 3}"] = "PL 頁延伸區塊；FY2030 起為收斂法"; T["eps_h"] = t0 + 3
    EPSR = {}
    for i, y in enumerate(YRS6):
        tset(t0 + 4 + i, f"eps{y}", f"  FY{y}", [f"=PL!{ext_cols[y]}{EXT[s]['eps']}" for s in SC], MF); EPSR[y] = t0 + 4 + i
    rr = t0 + 10
    V.cell(rr, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); V[f"E{rr}"] = "模型無折舊與資本支出項目，故不用（淨利＋折舊－資本支出）；設備股資本支出低（輕資產），此式偏誤小"; T["cf_h"] = rr
    tset(rr + 1, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
    RV = rr + 1; CFR = {}
    for i, y in enumerate(YRS6):
        tset(rr + 2 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], MF); CFR[y] = rr + 2 + i
    c1, c6 = CFR[2027], CFR[2032]
    tset(rr + 8, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], MF,
         "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027" + CFG["dcf_note"])
    tset(rr + 9, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{rr + 8}*(1+tp_r)" for c in "BCD"], MF)
    tset(rr + 10, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{rr + 9}+tp_wB*C{rr + 9}+tp_wC*D{rr + 9}", None, None], MF, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁 tp_wA／tp_wB／tp_wC")
    q = rr + 12
    V.cell(q, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; T["h02"] = q
    hdr(V, q + 1, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["phd"] = q + 1
    for i, (nm, tk, pe, dt, src) in enumerate(CFG["peers"], q + 2):
        V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
    pl_ = q + 1 + len(CFG["peers"]); T["p1"] = q + 2; T["pl"] = pl_
    q = pl_ + 1
    tset(q, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{T['p1']}:C{pl_})", None, None], "0.0", CFG["peer_note"])
    fw = CFG["fwd"]
    def feps(c): return "(" + "+".join(f"{w}*{c}{EPSR[y]}" for y, w in fw.items()) + ")"
    tset(q + 1, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=tp_wA*" + feps("B") + "+tp_wB*" + feps("C") + "+tp_wC*" + feps("D"), None, None], MF, CFG["fwd_note"])
    tset(q + 2, "comptp", "Comparable 目標價", [f"=B{q}*B{q + 1}", None, None], MF, "同業中位數 × 上列 EPS")
    q += 4
    V.cell(q, 1, "0.3 結論").font = BOLD; T["h03"] = q
    hdr(V, q + 1, ["項目", "數值", "", "", "說明"]); T["hd3"] = q + 1
    q += 2
    tset(q, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], MF, "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁")
    tset(q + 1, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], MF, "引用 0.2")
    tset(q + 2, "cons", "市場共識目標價", ["=cons_tp", None, None], MF, CFG["cons_tp"][1])
    tset(q + 3, "upw", "上檔空間：加權目標價", [f"=B{q}/px-1", None, None], "0.0%", CFG["px_note"])
    tset(q + 4, "upc", "上檔空間：Comparable 目標價", [f"=B{q + 1}/px-1", None, None], "0.0%")
    tset(q + 5, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型 12 個月後前瞻 EPS ÷ 共識前瞻 EPS − 1；" + CFG["cons_feps"][1])
    tset(q + 6, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%")
    V[f"E{q + 6}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
    tset(q + 7, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"{CFG["band_fmt"]}")&"～"&TEXT(B{T["dcf12"]},"{CFG["band_fmt"]}")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
    PE = CFG["pe10"]
    tset(q + 8, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
         [f"=IF(px/{c}{EPSR[2028]}<MIN({PE}),0,IF(px/{c}{EPSR[2028]}>MAX({PE}),1,PERCENTRANK({PE},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%", CFG["pe10_note"])
    V.cell(q + 9, 1, "目標價方法"); V[f"B{q + 9}"] = CFG["method"]; T["meth"] = q + 9
    for k in ("dcftp", "comptp", "tpw"): V[f"B{T[k]}"].fill = KEY
    if CFG.get("a3"): V["A3"] = CFG["a3"].format(row=t0)

    # --- Sens：目標價參數 ---
    s0 = SE.max_row + 3
    SE.cell(s0, 1, "目標價參數（v2 新增；只改 DCF 參數與權重，PL 收斂法的長期成長不隨 tp_g 變動）").font = BOLD
    hdr(SE, s0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
    def dcf_formula(c, r_="tp_r", g="tp_g"):
        cf = lambda y: f"Val!{c}{EPSR[y]}*(1-MIN({g}/tp_roe,0.8))"
        pv = "+".join(f"{cf(y)}/(1+{r_})^{t}" for t, y in enumerate(YRS6, 1))
        return f"({pv}+{cf(2032)}*(1+{g})/({r_}-{g})/(1+{r_})^6)"
    cases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r_="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r_="(tp_r-0.01)")),
             ("長期成長 +1pt", "g 5%（永續與再投資率）", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
             ("DCF 權重 0.4", "Comparable 0.6", dict(w="0.4")), ("DCF 權重 0.8", "Comparable 0.2", dict(w="0.8"))]
    for k, (nm, st, kw) in enumerate(cases):
        i = s0 + 2 + k; w = kw.pop("w", "tp_w_dcf")
        SE.cell(i, 1, nm); SE.cell(i, 2, st)
        rr_ = kw.get("r_", "tp_r")
        SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, **kw)}" for s, c in zip(SC, "BCD")) + f")*(1+{rr_})"
        SE[f"D{i}"] = f"=Val!B{T['comptp']}"
        SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
        SE[f"F{i}"] = f"=E{i}/E${s0 + 2}-1"
        for c in "CDE": SE[f"{c}{i}"].number_format = MF
        SE[f"F{i}"].number_format = "0.0%"
    T["sens0"] = s0; T["sens_last"] = s0 + 1 + len(cases)

    # --- README 頁 ---
    rr = RD.max_row + 1
    for t in CFG["readme"]:
        RD.cell(rr, 1, t); rr += 1
    return T, EXT

V2T, V2EXT = valuation_v2(CFG_V2)
if os.environ.get("V2_ROWS"):
    print("V2", V2T, V2EXT)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
