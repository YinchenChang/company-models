# Siliconomics（AI 半導體產業基準庫）建檔程式
# 用法：
#   首版（由種子建）：python3 siliconomics/builder/build.py --out <raw.xlsx>
#   之後各版（Excel 優先）：python3 siliconomics/builder/build.py --base siliconomics/model/<現行>.xlsx --out <raw.xlsx>
#     --base 時，In 頁的藍字輸入（低／基準／高、標記、來源）一律讀回現行 Excel；種子只用於 Excel 裡還沒有的鍵。
# 產出後以 LibreOffice headless 重算存檔（見 siliconomics/README.md），重算錯誤須為 0。
# 原則（比照 Tokenomics CLAUDE.md 第 1 節）：Excel 是唯一事實來源；本程式只產生公式頁與格式，不得在程式內新增或省略計算機制。
import argparse, json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
SIL = os.path.dirname(HERE)
ROOT = os.path.dirname(SIL)
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))

VERSION = "v0.1"
DATE = "20261027"
TK_SNAPSHOT = os.path.join(SIL, "data", "tokenomics_snapshot_v21.7.json")
TK_OLD = os.path.join(ROOT, "markets", "chip_bridge", "data", "tokenomics_snapshot_v20.1.json")
SEED = os.path.join(HERE, "seed", "seed_v0.1.json")
CONFLICTS = os.path.join(HERE, "seed", "conflicts_v0.1.json")
DECISIONS = os.path.join(HERE, "seed", "decisions_v0.1.json")
REF_SKH = os.path.join(ROOT, "skhynix", "dist")

BLUE = Font(color="0000FF"); GREEN = Font(color="008000"); BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7"); KEY = PatternFill("solid", fgColor="FFF2CC")
PEND = PatternFill("solid", fgColor="FFFF00"); WRAP = Alignment(wrap_text=True, vertical="top")
YEARS = [2025, 2026, 2027, 2028, 2029]


def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def load_base_inputs(path):
    """讀回現行 Excel 的 In 頁（Excel 優先）：{鍵: (低, 基準, 高, 標記, 來源)}"""
    if not path:
        return {}
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["In"]
    out = {}
    for r in range(5, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if not k or ws.cell(r, 3).value is None and ws.cell(r, 5).value is None:
            continue
        out[k] = tuple(ws.cell(r, c).value for c in (4, 5, 6, 7, 8))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    seed = json.load(open(SEED, encoding="utf-8"))
    base = load_base_inputs(a.base)
    tk = json.load(open(TK_SNAPSHOT, encoding="utf-8"))
    tk_old = json.load(open(TK_OLD, encoding="utf-8")) if os.path.exists(TK_OLD) else None
    conflicts = json.load(open(CONFLICTS, encoding="utf-8")) if os.path.exists(CONFLICTS) else []
    decisions = json.load(open(DECISIONS, encoding="utf-8")) if os.path.exists(DECISIONS) else []

    wb = openpyxl.Workbook()
    RD = wb.active; RD.title = "README"
    SC = wb.create_sheet("Scen")
    IN = wb.create_sheet("In")

    # ---------------- In：藍字輸入（Excel 擁有） ----------------
    IN["A1"] = "In — 產業輸入（藍字＝輸入，由本活頁簿擁有；要改值直接改本頁，builder 重建時讀回）。每列具名範圍＝鍵（指向 E 欄），另有 _lo／_hi（D、F 欄）"
    IN["A2"] = "標記：Verified／Interested-party（寫明是誰）／Analogy／Assumed（給低／高區間）／Derived（寫明換算）。v0.1 內容由 markets/memory/market.py 逐字搬入（數值、標記、來源文字不變）"
    hdr(IN, 4, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源／說明", "Interface 名稱"])
    INROW = {}
    r = 5
    for row in seed["rows"]:
        if "sec" in row:
            IN.cell(r, 1, row["sec"]).font = BOLD; r += 1; continue
        k = row["id"]
        lo, v, hi, tag, src = base.get(k, (row["lo"], row["base"], row["hi"], row["tag"], row["src"]))
        IN.cell(r, 1, k); IN.cell(r, 2, row["label"]); IN.cell(r, 3, row["unit"])
        for j, x in ((4, lo), (5, v), (6, hi)):
            if x is not None:
                c = IN.cell(r, j, x); c.font = BLUE
                if row["unit"] == "%": c.number_format = "0.0%"
        IN.cell(r, 7, tag); IN.cell(r, 8, src); IN.cell(r, 9, "SIF_" + k)
        name(wb, k, f"In!$E${r}")
        if lo is not None: name(wb, k + "_lo", f"In!$D${r}")
        if hi is not None: name(wb, k + "_hi", f"In!$F${r}")
        INROW[k] = (r, row["label"], row["unit"], lo is not None, hi is not None, tag)
        r += 1
    r += 1
    IN.cell(r, 1, "M2 分組：加速器平台（Units 頁各列與 market.py UNITS 的名稱、標記、來源）").font = BOLD; r += 1
    hdr(IN, r, ["分組鍵", "平台名稱", "", "", "", "", "標記", "來源／說明"]); r += 1
    GRP_ROW0 = r
    for g in seed["unit_groups"]:
        IN.cell(r, 1, "grp_" + g["key"]); IN.cell(r, 2, g["name"]); IN.cell(r, 7, g["tag"]); IN.cell(r, 8, g["src"])
        r += 1
    IN_LAST = r - 1
    for col, w in zip("ABCDEFGHI", (14, 52, 8, 9, 10, 9, 18, 70, 18)):
        IN.column_dimensions[col].width = w

    # ---------------- Path、Mkt：記憶體價格路徑與 HBM 市場（公式與 markets/memory/market.py write() 同一份） ----------------
    import market as MKT
    ref = MKT.write(wb, hdr, Font)
    MK = ref["mk"]
    wb["Mkt"]["A2"] = "公式由 markets/memory/market.py write() 產生（與各記憶體公司模型的 Mkt 頁同一份）；輸入讀本活頁簿 In 頁"

    # ---------------- Scen：情境定義 ----------------
    SC["A1"] = "Scen — 行業情境定義（A＝順利、B＝平平、C＝不如預期）。各公司模型的 A／B／C 一律讀同字母行業情境（Andy 2026-10-25 13:13）"
    hdr(SC, 3, ["情境", "名稱", "定義（顆數與部署）", "記憶體價格路徑", "2028 顆數 ÷ 基準", "2029 顆數 ÷ 基準", "2030–2032"])
    NM = {"A": "順利", "B": "平平", "C": "不如預期"}
    DEF = {"A": "基準顆數（市場層 2026–2028 預測）；部署順利、新技術導入快",
           "B": "AI 資本支出消化：2028 顆數約為基準 0.8、2029 約 0.85",
           "C": "明顯下修：2028、2029 顆數約為基準 0.6"}
    for i, s in enumerate(seed["scen"]):
        rr = 4 + i
        SC.cell(rr, 1, s); SC.cell(rr, 2, NM[s]); SC.cell(rr, 3, DEF[s]); SC.cell(rr, 4, seed["scen_desc"][s])
        for j, y in ((5, 28), (6, 29)):
            c = SC.cell(rr, j, f"=um{y}_{s}"); c.font = GREEN; c.number_format = "0.00"
        c = SC.cell(rr, 7, "未定：待算力收支 R4 產業部署路徑（預計 10-31 後）；各公司模型目前以收斂法自推"); c.fill = PEND
    SC["A9"] = "說明"
    SC["A9"].font = BOLD
    SC["A10"] = ("情境係數只定義 2028–29（Assumed，無外部預測可校準；B、C 區間 ±0.1，見 In 頁 M2b）。"
                 "R4 上線後，2026–2030 的部署 GW 三情境改讀算力收支（權威在算力收支，第 13 條），本頁的係數改為由部署路徑推導。")
    for col, w in zip("ABCDEFG", (6, 10, 46, 60, 14, 14, 50)):
        SC.column_dimensions[col].width = w

    # ---------------- TK_In：Tokenomics 鏡像（只讀） ----------------
    TKS = wb.create_sheet("TK_In")
    src = tk["source"]
    TKS["A1"] = (f"TK_In — Tokenomics 取數（只讀鏡像；權威在算力收支）：{src['file']}（{src['version']}，commit {src['commit'][:7]}，擷取 {src['extractedAt']}）。"
                 "綠字＝外部連結值，不得在本頁修改；Tokenomics 升版時以 tools/tokenomics/import_tokenomics.py 重抽快照再重建")
    TKS["A2"] = "右側 J–L 欄：company-models 下游目前讀的舊版快照（markets/chip_bridge v20.1）；M 欄＝與現行 TK 不同的格數（>0 表示下游用的值已過時）"
    hdr(TKS, 4, ["名稱（TK）", "世代／項目", "本庫名稱", "低", "基準", "高", "單位", "TK 位置", "取用者",
                 "舊快照 低", "舊快照 基準", "舊快照 高", "差異格數"])
    gens = {g["name"]: g["code"] for g in src["generations"]}
    users = {"IF_RacksPerGW": "chip_bridge、network（經橋接表）", "IF_GPUsPerGW": "chip_bridge", "IF_CapexIT": "chip_bridge、network、delta、quanta 等（部署 GW 換算）",
             "IF_CapexFacility": "chip_bridge", "IF_CapexTotal": "chip_bridge",
             "IF_RackPrice": "chip_bridge（目前抄錄值，改接排入中）", "IF_ScaleOutShare": "chip_bridge（目前抄錄值）", "IF_StoragePerGPU": "chip_bridge（目前抄錄值）"}
    r = 5; TKROWS = []
    old = tk_old["items"] if tk_old else {}
    for nm_, it in tk["items"].items():
        short = nm_[3:]
        if it.get("kind") == "gen_cost":
            for g, vv in it["values"].items():
                code = gens.get(g, g)
                sname = f"TK_{short}_{code}"
                TKS.cell(r, 1, nm_); TKS.cell(r, 2, g); TKS.cell(r, 3, sname)
                for j, c_ in ((4, "低成本"), (5, "基準"), (6, "高成本")):
                    TKS.cell(r, j, vv[c_]).font = GREEN
                TKS.cell(r, 7, it.get("unit")); TKS.cell(r, 8, it.get("ref")); TKS.cell(r, 9, users.get(nm_, "—（本庫已鏡像，待下游改接）"))
                if nm_ in old:
                    ov = old[nm_]["values"][g]
                    for j, c_ in ((10, "低成本"), (11, "基準"), (12, "高成本")):
                        TKS.cell(r, j, ov[c_])
                    TKS.cell(r, 13, f"=(ABS(D{r}-J{r})>1E-9)+(ABS(E{r}-K{r})>1E-9)+(ABS(F{r}-L{r})>1E-9)")
                name(wb, sname, f"TK_In!$E${r}")
                name(wb, sname + "_lo", f"TK_In!$D${r}"); name(wb, sname + "_hi", f"TK_In!$F${r}")
                TKROWS.append((sname, f"{it.get('label')}（{g}）", it.get("unit"), r)); r += 1
        else:
            vals = it.get("values") or it.get("value")
            sname = f"TK_{short}"
            TKS.cell(r, 1, nm_); TKS.cell(r, 2, it.get("label")); TKS.cell(r, 3, sname)
            if isinstance(vals, list):
                for j, x in zip((4, 5, 6), vals):
                    TKS.cell(r, j, x).font = GREEN
            else:
                TKS.cell(r, 5, vals).font = GREEN
            TKS.cell(r, 7, it.get("unit")); TKS.cell(r, 8, it.get("ref")); TKS.cell(r, 9, "avgo、mrvl、amd 等（目前各自假設，待改接）")
            name(wb, sname, f"TK_In!$E${r}")
            name(wb, sname + "_lo", f"TK_In!$D${r}"); name(wb, sname + "_hi", f"TK_In!$F${r}")
            TKROWS.append((sname, it.get("label"), it.get("unit"), r)); r += 1
    TK_LAST = r - 1
    r += 1
    TKS.cell(r, 1, "下游舊快照與現行 TK 不同的格數").font = BOLD
    c = TKS.cell(r, 13, f"=SUM(M5:M{TK_LAST})"); c.fill = KEY; TK_DIFF_CELL = f"TK_In!$M${r}"
    for col, w in zip("ABCDEFGHIJKLM", (22, 34, 26, 9, 9, 9, 10, 20, 34, 9, 9, 9, 8)):
        TKS.column_dimensions[col].width = w

    # ---------------- Conflicts：跨模型同一變數不同值的登錄 ----------------
    CF = wb.create_sheet("Conflicts")
    CF["A1"] = ("Conflicts — 跨模型衝突登錄：同一個行業變數在兩個以上公司模型有不同值。本庫的任務是讓每個變數只有一個值；"
                "建議值為 Claude 預設（判斷類，Andy 可改），在對應子產業併入本庫的版本生效，生效前各公司模型維持現值")
    hdr(CF, 3, ["ID", "子產業", "變數", "單位", "模型甲（位置）", "甲值", "模型乙（位置）", "乙值", "性質", "建議單一值與理由", "生效版本", "狀態", "來源"])
    r = 4
    for cf in conflicts:
        for j, k in enumerate(("id", "sub", "var", "unit", "a_loc", "a_val", "b_loc", "b_val", "kind", "proposal", "target", "status", "origin"), 1):
            c = CF.cell(r, j, cf.get(k)); c.alignment = WRAP
        if cf.get("status", "").startswith("待 Andy"):
            for j in range(1, 14): CF.cell(r, j).fill = PEND
        r += 1
    CF_LAST = max(r - 1, 4)
    for col, w in zip("ABCDEFGHIJKLM", (7, 10, 30, 9, 26, 12, 26, 12, 12, 52, 9, 14, 22)):
        CF.column_dimensions[col].width = w

    # ---------------- Interface：下游唯一連結點 ----------------
    IF = wb.create_sheet("Interface")
    IF["A1"] = ("Interface — 下游唯一連結點（各公司模型與市場層只讀本頁 SIF_ 具名範圍，經 tools/export_snapshot.py 匯成版本固定快照）。"
                "A 節＝輸入（純量，含低／高）；B 節＝推導（年度）；C 節＝Tokenomics 鏡像（SIF_TK_*，值＝TK_In）")
    IF["A2"] = f"版本：Siliconomics {VERSION}（{DATE}）；Tokenomics {src['version']}"
    hdr(IF, 4, ["名稱", "項目", "單位", "低", "基準", "高", "標記", "來源位置"])
    r = 5
    IF.cell(r, 1, "A. 輸入（純量）").font = BOLD; r += 1
    IFNAMES = []
    for k, (ir, lab, unit, has_lo, has_hi, tag) in INROW.items():
        nm_ = "SIF_" + k
        IF.cell(r, 1, nm_); IF.cell(r, 2, lab); IF.cell(r, 3, unit)
        if has_lo: IF.cell(r, 4, f"=In!D{ir}").font = GREEN
        c = IF.cell(r, 5, f"=In!E{ir}"); c.font = GREEN
        if has_hi: IF.cell(r, 6, f"=In!F{ir}").font = GREEN
        IF.cell(r, 7, f"=In!G{ir}"); IF.cell(r, 8, f"In!E{ir}")
        name(wb, nm_, f"Interface!$E${r}")
        if has_lo: name(wb, nm_ + "_lo", f"Interface!$D${r}")
        if has_hi: name(wb, nm_ + "_hi", f"Interface!$F${r}")
        IFNAMES.append(nm_); r += 1
    r += 1
    IF.cell(r, 1, "B. 推導（日曆年；欄 D–H＝2025–2029）").font = BOLD; r += 1
    hdr(IF, r, ["名稱", "項目", "單位", "2025", "2026", "2027", "2028", "2029", "來源位置"]); r += 1
    M = wb["Mkt"]
    derived = [("SIF_HBM_DemandEB", "dem", "HBM 需求（加速器裝載，情境 A 基準）", "EB"),
               ("SIF_HBM_DemandGrossEB", "demg", "HBM 需求 × 出貨係數（情境 A 基準）", "EB"),
               ("SIF_Units_Total", "units", "加速器顆數合計（情境 A 基準）", "百萬顆"),
               ("SIF_HBM_GBperUnit", "gbavg", "每顆平均 HBM", "GB"),
               ("SIF_HBM_SupplyEB", "sup", "HBM 位元供給（三情境相同）", "EB")]
    for s in seed["scen"]:
        derived += [(f"SIF_UnitsCoef_{s}", "cf" + s, f"顆數係數（{s}）", "x"),
                    (f"SIF_HBM_ShipEB_{s}", "ship" + s, f"HBM 出貨位元（{s}）", "EB"),
                    (f"SIF_HBM_GapEB_{s}", "gap" + s, f"HBM 需求 − 供給（{s}）", "EB"),
                    (f"SIF_HBM_Price_{s}", "p" + s, f"HBM 單價（{s}）", "$/GB"),
                    (f"SIF_HBM_Mkt_{s}", "m" + s, f"HBM 市場（{s}）", "$B")]
    for nm_, key, lab, unit in derived:
        mr = MK[key]
        IF.cell(r, 1, nm_); IF.cell(r, 2, lab); IF.cell(r, 3, unit)
        for j in range(5):
            c = IF.cell(r, 4 + j, f'=IF(ISBLANK(Mkt!{"CDEFG"[j]}{mr}),"",Mkt!{"CDEFG"[j]}{mr})'); c.font = GREEN; c.number_format = "0.00"
        IF.cell(r, 9, f"Mkt!C{mr}:G{mr}")
        name(wb, nm_, f"Interface!$D${r}:$H${r}")
        IFNAMES.append(nm_); r += 1
    r += 1
    IF.cell(r, 1, "C. Tokenomics 鏡像（只讀；權威在算力收支；本庫與 TK 一致由 Checks K3 檢查）").font = BOLD; r += 1
    hdr(IF, r, ["名稱", "項目", "單位", "低", "基準", "高", "", "來源位置"]); r += 1
    for sname, lab, unit, tr in TKROWS:
        nm_ = "SIF_" + sname
        IF.cell(r, 1, nm_); IF.cell(r, 2, lab); IF.cell(r, 3, unit)
        for j, col in ((4, "D"), (5, "E"), (6, "F")):
            IF.cell(r, j, f'=IF(ISBLANK(TK_In!{col}{tr}),"",TK_In!{col}{tr})').font = GREEN
        IF.cell(r, 8, f"TK_In!E{tr}")
        name(wb, nm_, f"Interface!$E${r}")
        name(wb, nm_ + "_lo", f"Interface!$D${r}"); name(wb, nm_ + "_hi", f"Interface!$F${r}")
        IFNAMES.append(nm_); r += 1
    for col, w in zip("ABCDEFGHI", (30, 50, 9, 10, 10, 10, 10, 14, 16)):
        IF.column_dimensions[col].width = w

    # ---------------- Checks ----------------
    CK = wb.create_sheet("Checks")
    CK["A1"] = "Checks — 交叉檢查與治理（ERROR 計入 SGOV_Errors，須為 0；WARN 計入 SGOV_Warn，列為待補）"
    hdr(CK, 3, ["ID", "項目", "本庫值", "參照值", "差", "判讀", "性質", "說明"])
    r = 4
    errs, warns = [], []
    # K1 搬移忠實度：HBM 市場與 SK 海力士成品 Mkt 頁逐值相同
    ref_vals = {}
    try:
        fn = [f for f in os.listdir(REF_SKH) if f.endswith(".xlsx")][0]
        rw = openpyxl.load_workbook(os.path.join(REF_SKH, fn), data_only=True)["Mkt"]
        for rr in range(1, rw.max_row + 1):
            lab = rw.cell(rr, 1).value
            if isinstance(lab, str) and lab.startswith("HBM 市場（"):
                ref_vals[lab] = [rw.cell(rr, c).value for c in range(3, 8)]
        ref_name = f"skhynix/dist/{fn}"
    except Exception as e:  # 找不到成品時只記錄
        ref_name = f"（讀取失敗：{e}）"
    for s in seed["scen"]:
        lab = f"HBM 市場（{s}）"
        for j, y in enumerate(YEARS):
            rv = ref_vals.get(lab, [None] * 5)[j]
            CK.cell(r, 1, f"K1-{s}{y}"); CK.cell(r, 2, f"{lab} {y}：本庫 vs SK 海力士模型成品")
            CK.cell(r, 3, f"=Mkt!{'CDEFG'[j]}{MK['m' + s]}")
            CK.cell(r, 4, rv)
            CK.cell(r, 5, f"=IF(ISNUMBER(D{r}),C{r}-D{r},\"\")")
            CK.cell(r, 6, f'=IF(NOT(ISNUMBER(D{r})),"WARN",IF(ABS(E{r})<1E-9,"OK","ERROR"))')
            CK.cell(r, 7, "驗證（搬移忠實度）"); CK.cell(r, 8, ref_name)
            r += 1
    # K2 情境係數
    for s in seed["scen"]:
        for y in (28, 29):
            CK.cell(r, 1, f"K2-{s}{y}"); CK.cell(r, 2, f"情境 {s} 20{y} 顆數係數：A＝1；B、C 在區間內")
            CK.cell(r, 3, f"=um{y}_{s}")
            if s == "A":
                CK.cell(r, 6, f'=IF(ABS(C{r}-1)<1E-9,"OK","ERROR")')
            else:
                CK.cell(r, 4, f'=um{y}_{s}_lo&"–"&um{y}_{s}_hi')
                CK.cell(r, 6, f'=IF(AND(C{r}>=um{y}_{s}_lo,C{r}<=um{y}_{s}_hi),"OK","ERROR")')
            CK.cell(r, 7, "驗證"); r += 1
        # 單調：A ≥ B ≥ C
    for y in (28, 29):
        CK.cell(r, 1, f"K2-ord{y}"); CK.cell(r, 2, f"20{y} 顆數係數 A ≥ B ≥ C")
        CK.cell(r, 6, f'=IF(AND(um{y}_A>=um{y}_B,um{y}_B>=um{y}_C),"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    # K3 TK 一致
    CK.cell(r, 1, "K3"); CK.cell(r, 2, "下游（chip_bridge 等）讀的 Tokenomics 快照與現行 TK 不同的格數")
    CK.cell(r, 3, f"={TK_DIFF_CELL}"); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理（TK 一致）")
    CK.cell(r, 8, "WARN＝下游仍用舊版 TK；本庫已用現行版；差異明細見 TK_In 頁 M 欄"); r += 1
    # G 治理：輸入列標記、來源、區間
    g0 = r
    CK.cell(r, 1, "G1"); CK.cell(r, 2, "In 頁輸入列缺標記")
    CK.cell(r, 3, f'=SUMPRODUCT((In!A5:A{IN_LAST}<>"")*(In!E5:E{IN_LAST}<>"")*(In!G5:G{IN_LAST}=""))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G2"); CK.cell(r, 2, "In 頁輸入列缺來源")
    CK.cell(r, 3, f'=SUMPRODUCT((In!A5:A{IN_LAST}<>"")*(In!E5:E{IN_LAST}<>"")*(In!H5:H{IN_LAST}=""))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理")
    CK.cell(r, 8, "v0.1 搬入時即存在（market.py 原狀：2029 HBM 單價三列來源欄空白）；與 G3 一併於 v0.2 補齊（屆時下游 In 頁文字會跟著變）"); r += 1
    n_noint = sum(1 for k, v in INROW.items() if ("Assumed" in str(v[5]) or "Analogy" in str(v[5])) and not (v[3] and v[4]))
    CK.cell(r, 1, "G3"); CK.cell(r, 2, "標記含 Assumed／Analogy 但未給低／高區間的輸入列（Tokenomics 規則：一律給區間）")
    CK.cell(r, 3, f'=SUMPRODUCT((In!A5:A{IN_LAST}<>"")*(In!E5:E{IN_LAST}<>"")*(ISNUMBER(SEARCH("Assumed",In!G5:G{IN_LAST}))+ISNUMBER(SEARCH("Analogy",In!G5:G{IN_LAST}))>0)*((In!D5:D{IN_LAST}="")+(In!F5:F{IN_LAST}="")>0))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理")
    CK.cell(r, 8, f"v0.1 搬入時即存在（market.py 原狀，建置時 {n_noint} 列）；區間待補，排入 v0.2"); r += 1
    keys = [k for k in INROW]
    CK.cell(r, 1, "G4"); CK.cell(r, 2, "In 頁鍵重複（建置時檢查）"); CK.cell(r, 3, len(keys) - len(set(keys)))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G5"); CK.cell(r, 2, "衝突登錄中『待 Andy』或『待併入』的項目數")
    CK.cell(r, 3, f'=COUNTIF(Conflicts!L4:L{CF_LAST},"待*")'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理")
    CK.cell(r, 8, "WARN＝仍有同一變數兩個值；各子產業併入本庫後歸零"); r += 1
    r += 1
    CK.cell(r, 1, "SGOV_Errors").font = BOLD; c = CK.cell(r, 3, f'=COUNTIF(F4:F{r-2},"ERROR")'); c.fill = KEY
    name(wb, "SGOV_Errors", f"Checks!$C${r}"); r += 1
    CK.cell(r, 1, "SGOV_Warn").font = BOLD; c = CK.cell(r, 3, f'=COUNTIF(F4:F{r-3},"WARN")'); c.fill = KEY
    name(wb, "SGOV_Warn", f"Checks!$C${r}")
    for col, w in zip("ABCDEFGH", (10, 58, 12, 12, 10, 8, 18, 60)):
        CK.column_dimensions[col].width = w

    # ---------------- Decisions、DB_Evidence ----------------
    DC = wb.create_sheet("Decisions")
    DC["A1"] = "Decisions — 決策登錄（Excel 擁有）。判斷類由 Andy 決定；Claude 依預設執行者標『預設、可撤回』"
    hdr(DC, 3, ["決策 ID", "類別", "項目", "決議", "日期", "Andy 原話／依據", "狀態"])
    r = 4
    for d in decisions:
        for j, k in enumerate(("id", "cat", "item", "decision", "date", "quote", "status"), 1):
            c = DC.cell(r, j, d.get(k)); c.alignment = WRAP
        if "預設" in str(d.get("status")):
            for j in range(1, 8): DC.cell(r, j).fill = PEND
        r += 1
    for col, w in zip("ABCDEFG", (9, 12, 24, 60, 12, 50, 16)):
        DC.column_dimensions[col].width = w
    EV = wb.create_sheet("DB_Evidence")
    EV["A1"] = "DB_Evidence — 證據登錄（所有比對過的新資訊，含不採納者）；先與本庫現值比對，再判定採納／部分採納／不採納／待查"
    hdr(EV, 3, ["ID", "日期", "主張（摘要）", "來源", "標記", "對應鍵", "本庫當時值", "新資訊值", "判定", "處理版本", "立場", "備註"])
    EV.cell(4, 1, "（v0.1 尚無；跨模型衝突先登錄於 Conflicts 頁，判定後再轉入本頁）")

    # ---------------- README ----------------
    lines = [
        (f"Siliconomics {VERSION} — AI 半導體產業基準庫（{DATE}）", True),
        ("用途：AI 半導體各子產業的『唯一事實』——加速器顆數、HBM 與記憶體價格、網通內容、CoWoS、載板、伺服器組裝、電源散熱、設備等行業變數只有一個值；各公司模型只讀本庫 Interface 的 SIF_ 名稱。", False),
        ("與 Tokenomics 的分工：每 GW 硬體量、機架價、每 GW 資本支出、NonNV 比例、產業年度部署路徑（R4）權威在 Tokenomics／算力收支，本庫只在 TK_In 頁鏡像（綠字、不得修改）；晶片產業資訊（HBM 單價與容量、每顆內容、網通內容、份額、出貨顆數、ASIC 規格、晶圓與設備）權威在本庫（《情境卡產生規則》第 13 條）。", False),
        ("v0.1 範圍：(1) Tokenomics v21.7 鏡像（22 個名稱）與下游舊快照差異；(2) 加速器顆數層、記憶體價格路徑、HBM 市場（由 markets/memory/market.py 逐字搬入，數值不變）；(3) 跨模型衝突登錄；(4) Interface、Checks、Decisions。網通、CoWoS、載板、組裝電源散熱、設備依規劃書分段併入。", False),
        ("讀法：Scen（情境定義）→ In（藍字輸入）→ Path／Mkt（推導）→ Interface（下游取數）→ Checks（SGOV_Errors 須為 0）。", False),
        ("顏色：藍字＝輸入（本活頁簿擁有）；黑字＝公式；綠字＝跨頁或外部連結；淡黃底＝關鍵輸出；亮黃底＝待 Andy 決定或未定。", False),
        ("標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy 與 Assumed 一律給區間（v0.1 搬入列尚有缺區間者，見 Checks G3）。", False),
        ("Excel 優先：要改輸入，直接改 In 頁藍字；builder 重建時讀回（--base）。公式頁不要手改（重建時會被覆寫）。", False),
        ("取數規則（下游）：只取 SIF_ 具名範圍（含 SIF_TK_* 鏡像）；不得連結 In、Mkt 等內部頁；以 tools/export_snapshot.py 匯成版本固定快照，版本號寫入公司模型。", False),
        ("版本命名：YYYYMMDD_Siliconomics_vX.Y.xlsx；model/CURRENT 記錄現行檔名。只改證據或標記為修補版；輸入值或公式改變為次版；Interface 名稱增刪為主版。", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = RD.cell(i, 1, t); c.alignment = WRAP
        if b: c.font = Font(bold=True, size=13)
    RD.column_dimensions["A"].width = 150

    wb.save(a.out)
    print(f"寫出 {a.out}：In {len(INROW)} 鍵、Interface {len(IFNAMES)} 名稱、衝突 {len(conflicts)} 項")


if __name__ == "__main__":
    main()
