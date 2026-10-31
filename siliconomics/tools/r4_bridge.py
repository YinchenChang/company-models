# R4 對帳橋接（Siliconomics 側）：Siliconomics 部署 GW（成本等值、出貨年）→ 推得「淨新增上線 IT GW」，與 Wattonomics 主路徑並列，殘差歸 Wattonomics 側。
# 用法：python3 siliconomics/tools/r4_bridge.py <輸出 xlsx>
#   讀 siliconomics/model/CURRENT 與 wattonomics/model/CURRENT 的重算值（data_only），寫成「輸入格（藍字）＋公式格（黑字）」的活頁簿；
#   之後以 LibreOffice 重算即可看到結果。本工具不改任何模型。
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SIL = os.path.join(ROOT, "siliconomics", "model")
WT = os.path.join(ROOT, "wattonomics", "model")
sil_f = open(os.path.join(SIL, "CURRENT")).read().strip()
wt_f = open(os.path.join(WT, "CURRENT")).read().strip()
S = openpyxl.load_workbook(os.path.join(SIL, sil_f), data_only=True)
W = openpyxl.load_workbook(os.path.join(WT, wt_f), data_only=True)


def nv(wb, name):
    """讀具名範圍的單一值。"""
    d = wb.defined_names[name]
    sh, ref = list(d.destinations)[0]
    return wb[sh][ref.replace("$", "")].value


def tk(name):
    """讀 Siliconomics TK_In 頁（Tokenomics 鏡像）某名稱的 低／基準／高。"""
    ws = S["TK_In"]
    for r in ws.iter_rows(values_only=True):
        if r[2] == name:
            return r[3], r[4], r[5]
    raise KeyError(name)


YRS = [2025, 2026, 2027, 2028, 2029]
YC = {y: 4 + i for i, y in enumerate(YRS)}  # 欄 D..H
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")
wb = openpyxl.Workbook()

# ---------------- In 頁：輸入 ----------------
IN = wb.active
IN.title = "In"
IN["A1"] = f"In — R4 對帳橋接的輸入。藍字＝輸入（取自 {sil_f}、{wt_f} 或本次查得的外部說法）；低／高為敏感度。標記沿用專案規則。"
IN["A1"].font = BOLD
hdr = ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源"]
for j, h in enumerate(hdr, 1):
    c = IN.cell(3, j, h); c.font = BOLD; c.fill = HDR
names = {}
r = 4


def inp(key, label, unit, lo, base, hi, tag, src):
    global r
    vals = [key, label, unit, lo, base, hi, tag, src]
    for j, v in enumerate(vals, 1):
        c = IN.cell(r, j, v)
        if j in (4, 5, 6) and v is not None:
            c.font = BLUE
    names[key] = f"In!$E${r}"
    names[key + "_lo"] = f"In!$D${r}"
    names[key + "_hi"] = f"In!$F${r}"
    r += 1


# A. Siliconomics 平台別 AI IT 資本支出（$B；情境 A＝基準）
DEP = S["Deploy"]
UN = S["Units"]
it = {k: nv(S, n) for k, n in [("GB", "it_GB"), ("VR", "it_VR"), ("RU", "it_RU"), ("HP", "it_HP"), ("AMD8", "it_AMD8"),
                                ("AMDR", "it_AMDR"), ("TPU", "it_TPU"), ("TRN", "it_TRN"), ("ASIC", "it_ASIC")]}
IN.cell(r, 1, "A. Siliconomics 平台別 AI IT 資本支出（百萬顆 × 每顆 IT capex 千美元 ＝ $B；情境 A，未乘情境係數）").font = BOLD; r += 1
plat = [("bw", 4, lambda y: it["GB"]), ("rb", 5, lambda y: it["VR"]), ("ru", 6, lambda y: it["RU"]), ("hp", 7, lambda y: it["HP"]),
        ("amd", 8, lambda y: it["AMD8"] if y == 2026 else it["AMDR"]), ("tpu", 9, lambda y: it["TPU"]),
        ("trn", 10, lambda y: it["TRN"]), ("oth", 11, lambda y: it["ASIC"])]
for k, row, fit in plat:
    for y in (2026, 2027, 2028, 2029):
        u = UN.cell(row, YC[y] - 2).value  # Units 頁 C..F＝2026..2029
        inp(f"cap_{k}_{y}", f"{UN.cell(row, 1).value} AI IT capex {y}", "$B", None, u * fit(y), None, "Derived",
            f"{sil_f} Units!{CL(YC[y]-2)}{row}（{u} 百萬顆）× 每顆 IT capex {fit(y):.2f} 千美元（In 頁 M9）")
for y in (2026, 2027, 2028, 2029):
    inp(f"dep_{y}", f"Siliconomics 部署 GW（成本等值，情境 A）{y}", "GW", None, DEP.cell(9, YC[y] - 2).value, None, "Derived",
        f"{sil_f} Deploy!{CL(YC[y]-2)}9＝AI IT capex ÷ 每 GW IT capex")
    inp(f"pgw_{y}", f"Siliconomics 每 GW IT capex {y}（成本等值的分母）", "$B/GW", None, DEP.cell(8, YC[y] - 2).value, None, "Derived",
        f"{sil_f} Deploy!{CL(YC[y]-2)}8（2026 依 vr26 混合 GB300／VR200；2027 起 VR200）")

# B. 非 NVIDIA 實體 GW 換算（Tokenomics 持有成本比）
IN.cell(r, 1, "B. 非 NVIDIA 實體 GW 換算：每 GW IT capex ＝ VR200 每 GW IT capex × 持有成本比（與 Wattonomics CapexGW 同法）").font = BOLD; r += 1
vr = tk("TK_CapexIT_VR200")
inp("pgw_vr", "VR200 每 IT GW IT capex", "$B/GW", vr[0], vr[1], vr[2], "Derived", "Tokenomics v21.7 IF_CapexIT（TK_In 鏡像）")
for k, nm, lab in [("amd", "TK_NonNV_Hold_AMD", "AMD（MI455X）"), ("tpu", "TK_NonNV_Hold_TPU", "Google TPU（Ironwood）"),
                   ("trn", "TK_NonNV_Hold_Trn", "AWS Trainium3")]:
    v = tk(nm)
    inp(f"hold_{k}", f"持有成本比 {lab} ÷ VR200", "x", v[0], v[1], v[2], "Assumed", f"Tokenomics v21.7 {nm.replace('TK_', 'IF_')}（TK 標 Assumed）")
v1, v2, v3 = tk("TK_NonNV_Hold_Maia"), tk("TK_NonNV_Hold_Ascend"), tk("TK_NonNV_Hold_OAI")
inp("hold_oth", "持有成本比 其他 ASIC ÷ VR200（Maia、昇騰、OpenAI 客製三者簡單平均）", "x",
    round((v1[0] + v2[0] + v3[0]) / 3, 3), round((v1[1] + v2[1] + v3[1]) / 3, 3), round((v1[2] + v2[2] + v3[2]) / 3, 3), "Assumed",
    f"Tokenomics v21.7 IF_NonNV_Hold_Maia {v1[1]}、_Ascend {v2[1]}、_OAI {v3[1]}；MTIA 無 TK 值，組合權重未知故取簡單平均")

# C. 出貨 → 上線
IN.cell(r, 1, "C. 出貨年 → 上線年").font = BOLD; r += 1
inp("lag", "出貨到上線的平均延遲（年）", "年", 0.25, 0.375, 0.5, "Assumed",
    "加速器出貨後經 ODM 整機、運送、上架、驗收到供電運轉；取 3–6 個月，基準 4.5 個月。來源：未找到可量化的公開統計")
inp("nv_dc_25", "NVIDIA 資料中心營收 FY2026（約 CY2025）", "$B", None, 193.7, None, "Verified",
    "NVIDIA 8-K：FY26 Q1–Q4 39.1／41.1／51.2／62.3（nvidia/build_nvidia.py Q 頁）")
inp("nv_dc_26", "NVIDIA 資料中心營收 FY2027 錨點（約 CY2026）", "$B", None, 375.46, None, "Verified／Assumed",
    "Q1、Q2 實績 75.2／89.0（8-K）＋Q3 指引中點＋Q4 假設（NVIDIA 個股模型 v0.2.3）")
inp("pgw_25", "每 GW IT capex 2025（NVIDIA 世代加權）", "$B/GW", None, 26.01, None, "Derived",
    f"{wt_f} CapexGW「NVIDIA 每 IT GW IT capex（世代加權）」2025 欄")
inp("nonnv_25", "非 NVIDIA 2025 實體 GW ÷ 2026", "x", 0.5, 0.6, 0.75, "Assumed",
    "2025 TPU、Trainium、AMD 顆數本庫未建；TPU v6→v7、Trn2→Trn3 放量，取約六成")

# D. 未上線庫存、汰換
IN.cell(r, 1, "D. 通路與未上線庫存、汰換（毛→淨）").font = BOLD; r += 1
inp("inv_k", "已出貨未上線庫存 ÷ 年出貨（庫存增加＝k × 出貨年增量）", "x", 0.05, 0.1, 0.2, "Assumed",
    "約 5 週出貨量（基準）；2026 年雲端客戶談到『有晶片沒電』（下方 E 節 Gerstner 說法）屬高值方向")
inp("retire", "汰換率（舊加速器下線釋出的電力 ÷ 前一年年底在役 IT GW）", "%", 0.0, 0.02, 0.05, "Assumed",
    "A100／H100 折舊 5–6 年，2026–27 下線少；高值反映 Hopper 換 Blackwell 就地更新")

# E. NVIDIA 每 GW 內容交叉檢核（只作敏感度）
IN.cell(r, 1, "E. NVIDIA 每 GW 內容交叉檢核（只作敏感度，不入基準）").font = BOLD; r += 1
inp("nvc_gb_tk", "Tokenomics 隱含 NVIDIA 內容 每 IT GW（GB300）", "$B/GW", None, round(487008 * nv(S, "nv_content") / 1e6, 2), None, "Derived",
    f"TK IF_GPUsPerGW GB300 487,008 × 每顆 NVIDIA 內容 {nv(S, 'nv_content')} 千美元（Siliconomics nv_content）")
inp("nvc_gb_jh", "黃仁勳說法：NVIDIA 每 GW 內容（Grace Blackwell）", "$B/GW", None, 25, None, "Interested-party",
    "247wallst 2026-09-24 轉述：Hopper 約 18、Grace Blackwell 約 25、Vera Rubin 約 40（十億美元／GW）；Bg2 Pod 2026-10-15「18 → 20 some → 40 plus」。GW 口徑（IT 或廠區）未說明")
inp("nvc_vr_jh", "黃仁勳說法：NVIDIA 每 GW 內容（Vera Rubin）", "$B/GW", None, 40, None, "Interested-party", "同上")
inp("gerst_27", "Gerstner 說法：2027 年新增 GW", "GW", None, 43, None, "Interested-party",
    "TIKR 轉述 Bg2 Pod 2026-10-15；同集 Tang：已宣布晶片約 50 GW、土地電力廠房 35–40 GW")

# F. Wattonomics 對照
IN.cell(r, 1, "F. Wattonomics 對照（主路徑、需求情景 A、D0）").font = BOLD; r += 1
CMP = W["Compare"]
DPL = W["Deploy"]


def wt_row(ws, label, scen, delay):
    for rr in ws.iter_rows(values_only=True):
        if rr[0] == label and rr[1] == scen and rr[2] == delay:
            return rr
    raise KeyError(label)


new = wt_row(DPL, "年新增 IT GW（淨，不含新納入存量）", "A", "D0")
ye = wt_row(DPL, "年底上線 IT GW", "A", "D0")
site = wt_row(CMP, "瓦帳 年新增 IT GW（淨；SiteOnly）", "A", "D0")
WTY = {2023: 3, 2024: 4, 2025: 5, 2026: 6, 2027: 7, 2028: 8, 2029: 9}
for y in YRS:
    inp(f"wt_new_{y}", f"Wattonomics 年新增 IT GW（主路徑）{y}", "GW", None, new[WTY[y]], None, "Derived", f"{wt_f} Deploy（A、D0）")
    inp(f"wt_site_{y}", f"Wattonomics 年新增 IT GW（SiteOnly）{y}", "GW", None, site[WTY[y]], None, "Derived", f"{wt_f} Compare（A、D0）")
    inp(f"wt_ye_{y-1}", f"Wattonomics 年底上線 IT GW {y-1}", "GW", None, ye[WTY[y - 1]], None, "Derived", f"{wt_f} Deploy（A、D0）")

for j, w in enumerate([16, 58, 8, 9, 10, 9, 16, 90], 1):
    IN.column_dimensions[CL(j)].width = w

for k, ref in names.items():
    wb.defined_names[k] = openpyxl.workbook.defined_name.DefinedName(k, attr_text=ref)

# ---------------- Bridge 頁 ----------------
B = wb.create_sheet("Bridge")
B["A1"] = ("Bridge — Siliconomics 部署 GW（成本等值、出貨年）→ Siliconomics 推得的淨新增上線 IT GW，與 Wattonomics 主路徑並列。"
           "殘差＝Wattonomics 側尚未解釋的部分（覆蓋率；算力收支補量化）。情境 A；B、C 的比值與 A 相同，因 Wattonomics 2028–29 目前讀 SIF_UnitsCoef")
B["A1"].font = BOLD
for j, h in enumerate(["列", "項目", "說明"] + [str(y) for y in YRS], 1):
    c = B.cell(3, j, h); c.font = BOLD; c.fill = HDR
R = {}
br = 4


def brow(key, label, note, fmap, bold=False):
    global br
    B.cell(br, 1, key); B.cell(br, 2, label); B.cell(br, 3, note)
    for y in YRS:
        f = fmap(y)
        if f is not None:
            B.cell(br, YC[y], f).number_format = "0.00"
    if bold:
        for j in range(1, 9):
            B.cell(br, j).font = BOLD
    R[key] = br
    br += 1


C = lambda y: CL(YC[y])
NONNV = [("amd", "hold_amd"), ("tpu", "hold_tpu"), ("trn", "hold_trn"), ("oth", "hold_oth")]
brow("b1", "Siliconomics 部署 GW（成本等值；出貨年）", "Deploy 第 9 列", lambda y: None if y == 2025 else f"=dep_{y}", True)
brow("b2a", "　非 NVIDIA 成本等值 GW", "Σ 非 NVIDIA capex ÷ 每 GW IT capex（本庫分母）",
     lambda y: None if y == 2025 else "=(" + "+".join(f"cap_{k}_{y}" for k, _ in NONNV) + f")/pgw_{y}")
brow("b2b", "　非 NVIDIA 實體 GW", "Σ capex ÷（VR200 每 GW × 持有成本比）",
     lambda y: None if y == 2025 else "=" + "+".join(f"cap_{k}_{y}/(pgw_vr*{h})" for k, h in NONNV))
brow("b2", "① 成本等值 → 實體（ASIC、AMD 每 GW 較便宜）", "b2b − b2a；方向：放大差距", lambda y: None if y == 2025 else f"={C(y)}{R['b2b']}-{C(y)}{R['b2a']}")
brow("b3", "實體出貨 GW", "b1 + ①；2025＝NVIDIA 營收比 × 世代每 GW 成本換算＋非 NVIDIA 比例",
     lambda y: (f"=({C(2026)}{R['b1']}-{C(2026)}{R['b2a']})*(nv_dc_25/nv_dc_26)*(pgw_2026/pgw_25)+{C(2026)}{R['b2b']}*nonnv_25"
                if y == 2025 else f"={C(y)}{R['b1']}+{C(y)}{R['b2']}"), True)
brow("b4", "② 出貨年 → 上線年", "−lag ×（當年實體出貨 − 前一年）",
     lambda y: None if y == 2025 else f"=-lag*({C(y)}{R['b3']}-{CL(YC[y]-1)}{R['b3']})")
brow("b5", "③ 已出貨未上線庫存增加", "−inv_k ×（當年實體出貨 − 前一年）",
     lambda y: None if y == 2025 else f"=-inv_k*({C(y)}{R['b3']}-{CL(YC[y]-1)}{R['b3']})")
brow("b6", "④ 汰換（毛 → 淨）", "−汰換率 × Wattonomics 前一年年底在役 GW",
     lambda y: None if y == 2025 else f"=-retire*wt_ye_{y-1}")
brow("b7", "⑤ 中國", "兩邊都含中國（Siliconomics：H200 等出貨；Wattonomics：Rystad 容量 × AI 占比），不調整；見 b7m",
     lambda y: None if y == 2025 else 0)
brow("b7m", "　（揭露）Siliconomics 中國 H200 成本等值 GW", "hp capex ÷ 每 GW；昇騰在『其他 ASIC』內未拆",
     lambda y: None if y == 2025 else f"=cap_hp_{y}/pgw_{y}")
brow("b8", "Siliconomics 推得 淨新增上線 IT GW", "b3 + ② + ③ + ④ + ⑤",
     lambda y: None if y == 2025 else f"={C(y)}{R['b3']}+{C(y)}{R['b4']}+{C(y)}{R['b5']}+{C(y)}{R['b6']}+{C(y)}{R['b7']}", True)
brow("b9", "Wattonomics 年新增 IT GW（主路徑，A、D0）", "", lambda y: f"=wt_new_{y}", True)
brow("b9s", "　Wattonomics SiteOnly", "只有站點證據", lambda y: f"=wt_site_{y}")
brow("b10", "⑥ 殘差（Wattonomics 側未解釋；歸覆蓋率）", "b8 − b9；由算力收支以 Wattonomics lag 與 C0 園區列量化（11-07）",
     lambda y: None if y == 2025 else f"={C(y)}{R['b8']}-{C(y)}{R['b9']}", True)
brow("b11", "比值：Wattonomics ÷ Siliconomics 部署 GW（原始）", "Compare 頁同值", lambda y: None if y == 2025 else f"={C(y)}{R['b9']}/{C(y)}{R['b1']}")
brow("b12", "比值：Wattonomics ÷ Siliconomics 推得上線 GW（橋接後）", "", lambda y: None if y == 2025 else f"={C(y)}{R['b9']}/{C(y)}{R['b8']}", True)
brow("b13", "殘差占 Siliconomics 部署 GW", "", lambda y: None if y == 2025 else f"={C(y)}{R['b10']}/{C(y)}{R['b1']}")
br += 1
B.cell(br, 1, "敏感度（2026、2027 推得上線 GW；一次動一個輸入，其餘取基準）").font = BOLD; br += 1
for j, h in enumerate(["列", "情況", "說明", "2026 推得上線", "2027 推得上線", "2026 殘差", "2027 殘差"], 1):
    c = B.cell(br, j, h); c.font = BOLD; c.fill = HDR
br += 1


def derived(y, lag="lag", inv="inv_k", ret="retire", hold=None, non25="nonnv_25", nvs="1"):
    """以文字公式重算 b8（年 y），可替換某一輸入。"""
    hmap = dict(NONNV)
    if hold:
        hmap = {k: h + hold for k, h in NONNV}
    def b2b(yy):
        return "(" + "+".join(f"cap_{k}_{yy}/(pgw_vr*{hmap[k]})" for k, _ in NONNV) + ")"
    def b2a(yy):
        return "((" + "+".join(f"cap_{k}_{yy}" for k, _ in NONNV) + f")/pgw_{yy})"
    def b3(yy):
        if yy == 2025:
            return f"((dep_2026-{b2a(2026)})*{nvs}*(nv_dc_25/nv_dc_26)*(pgw_2026/pgw_25)+{b2b(2026)}*{non25})"
        return f"((dep_{yy}-{b2a(yy)})*{nvs}+{b2b(yy)})"
    return f"{b3(y)}-({lag}+{inv})*({b3(y)}-{b3(y-1)})-{ret}*wt_ye_{y-1}"


cases = [("基準", {}), ("延遲 低（3 個月）", dict(lag="lag_lo")), ("延遲 高（6 個月）", dict(lag="lag_hi")),
         ("庫存 低", dict(inv="inv_k_lo")), ("庫存 高", dict(inv="inv_k_hi")), ("汰換 0", dict(ret="retire_lo")), ("汰換 高", dict(ret="retire_hi")),
         ("持有成本比 低（非 NVIDIA 更便宜→實體 GW 更多）", dict(hold="_lo")), ("持有成本比 高", dict(hold="_hi")),
         ("2025 非 NVIDIA 低", dict(non25="nonnv_25_lo")), ("2025 非 NVIDIA 高", dict(non25="nonnv_25_hi"))]
for nm, kw in cases:
    B.cell(br, 2, nm)
    B.cell(br, 4, "=" + derived(2026, **kw)).number_format = "0.00"
    B.cell(br, 5, "=" + derived(2027, **kw)).number_format = "0.00"
    B.cell(br, 6, f"=D{br}-wt_new_2026").number_format = "0.00"
    B.cell(br, 7, f"=E{br}-wt_new_2027").number_format = "0.00"
    br += 1
# NVIDIA 每 GW 內容檢核：若黃仁勳 GB 口徑正確，NVIDIA 實體 GW ＝ TK 的 (TK 內容 ÷ 25) 倍
B.cell(br, 2, "NVIDIA 每 GW 內容取黃仁勳 GB 說法（NVIDIA 實體 GW × TK 內容 ÷ 25）")
B.cell(br, 3, "Interested-party；GW 口徑未說明，只作方向")
B.cell(br, 4, "=" + derived(2026, nvs="(nvc_gb_tk/nvc_gb_jh)")).number_format = "0.00"
B.cell(br, 5, "=" + derived(2027, nvs="(nvc_gb_tk/nvc_gb_jh)")).number_format = "0.00"
B.cell(br, 6, f"=D{br}-wt_new_2026").number_format = "0.00"
B.cell(br, 7, f"=E{br}-wt_new_2027").number_format = "0.00"
br += 1
br += 2
notes = [
    "讀法：① 使差距變大（ASIC 每 GW 較便宜，同樣金額買到更多 GW）；②③④ 使差距縮小；⑤ 兩邊口徑相同、不調整。",
    "殘差主要是 Wattonomics 2026 年主路徑＝SiteOnly（只有站點證據；大型雲端公司多不公布遠期園區 MW），2027 起七家改用資本支出推得（甲2），其餘玩家仍只有站點證據。",
    "改讀設計（差距可解釋後）：Siliconomics 讀 Wattonomics 上線 GW，轉成出貨年＝(1−lag)×上線(t)＋lag×上線(t＋1)，再加回庫存增加與汰換，最後扣 ① 回到成本等值口徑，才與現行市場公式相容。",
]
for t in notes:
    B.cell(br, 1, t); br += 1
B.column_dimensions["A"].width = 6; B.column_dimensions["B"].width = 50; B.column_dimensions["C"].width = 60
for y in YRS:
    B.column_dimensions[C(y)].width = 11

out = sys.argv[1]
wb.save(out)
print("wrote", out)
