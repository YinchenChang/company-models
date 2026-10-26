"""AI 投資地圖 — 由資料活頁簿（LibreOffice 重算後的值）產生網站 dist/index.html。

網站只呈現活頁簿的值，不做任何計算（排序、篩選、繪圖除外）。
用法：python3 markets/invest_map/build_html.py
"""
import csv
import glob
import json
import re
import os
import shutil
import subprocess
import tempfile

import openpyxl

from profiles_md import load_all

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))


def recalc(path):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, path],
                   check=True, capture_output=True)
    out = os.path.join(tmp, os.path.basename(path))
    return out, tmp


def rows(ws):
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        yield dict(zip(hdr, r))


def num(x):
    return x if isinstance(x, (int, float)) else None


def main():
    books = sorted(glob.glob(os.path.join(HERE, "dist", "*_AI投資地圖_資料_*.xlsx")))
    src = books[-1]
    calc, tmp = recalc(src)
    wb = openpyxl.load_workbook(calc, data_only=True)
    shutil.rmtree(tmp, ignore_errors=True)

    errors = []
    for ws in wb:
        for r in ws.iter_rows(values_only=True):
            for v in r:
                if isinstance(v, str) and (v.startswith("#") or v.startswith("Err:")):
                    errors.append((ws.title, v))
    if errors:
        raise SystemExit(f"活頁簿有錯誤值：{errors[:5]}")

    snap = json.load(open(sorted(glob.glob(os.path.join(HERE, "data", "tokenomics_snapshot_v*.json")))[-1],
                          encoding="utf-8"))["source"]
    cm = subprocess.run(["git", "-C", ROOT, "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()

    pergw = {}
    for r in wb["每GW"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            pergw[f"{r[0]}｜{r[1]}"] = {"lo": num(r[2]), "base": num(r[3]), "hi": num(r[4]), "unit": r[5],
                                         "name": r[6], "cell": r[7]}

    segments = [{"layer": d["層級"], "name": d["次產業"], "share": d["占 AI IT 資本支出"],
                 "mid": num(d["比例中值"]), "usd": num(d["每 GW 金額（GB300，$B）"]),
                 "irr": d["不可替代性（1–3）"], "vc": d["價值鏈分量分"], "basis": d["依據"]}
                for d in rows(wb["價值鏈"])]

    position = [{"tk": str(d["代號"]), "name": d["公司"], "layer": d["層級"], "seg": d["次產業"],
                 "imp": d["公司重要性（1–3）"], "posn": d["市場位置（1–4）"], "basis": d["地位依據"], "tag": d["標記"],
                 "irr_s": d["不可替代性分"], "pos_s": d["市場位置分"], "vc_s": d["價值鏈分量分"],
                 "score": d["地位分數"], "rank": d["排名"]} for d in rows(wb["地位"])]

    prices = {str(d["代號"]): d for d in rows(wb["價格"])}
    sources = {}
    with open(os.path.join(HERE, "data", "prices.csv"), encoding="utf-8") as f:
        for d in csv.DictReader(f):
            sources[d["ticker"]] = d

    mx = json.load(open(os.path.join(HERE, "data", "models_extract.json"), encoding="utf-8"))
    MFOLDER, MFILE = {}, {}
    for tk_, v_ in mx.items():
        folder = v_["model"].split("/")[0] if "/" in v_["model"] else None
        if folder and glob.glob(os.path.join(ROOT, folder, "dist", "*.html")):
            MFOLDER[tk_] = folder
            MFILE[tk_] = os.path.basename(glob.glob(os.path.join(ROOT, folder, "dist", "*.html"))[0])
        else:
            MFILE[tk_] = os.path.basename(v_["model"])
    valuation = []
    for d in rows(wb["估值"]):
        tk = str(d["代號"])
        p = prices[tk]
        valuation.append({
            "tk": tk, "name": d["公司"], "layer": d["層級"], "seg": d["次產業"], "method": d["方法"],
            "ccy": p["幣別"], "price": num(d["採用價"]), "price_date": sources.get(tk, {}).get("date", p["模型價日期"]),
            "price_src": sources.get(tk, {}).get("source", ""), "price_n": sources.get(tk, {}).get("sources_count", ""),
            "model_px": num(p["模型內現價"]), "model_date": p["模型價日期"], "move": num(d["模型價日以來漲跌"]),
            "flag": d["價格變動旗標"] or "",
            "center": num(d["中心價"]), "bear": num(d["悲觀情境價"]), "cu": num(d["中心空間"]), "bu": num(d["悲觀情境空間"]),
            "tier": d["最終分級"], "vtier": d["分級"],
            "mfolder": MFOLDER.get(tk), "mfile": MFILE.get(tk, ""),
            "gate": d["收支關卡"], "gbasis": d["收支口徑"], "glabel": d["收支列（模型）"] or "", "gseries": d["收支序列"] or "",
            "gsum": num(d["期間合計"]), "gmin": num(d["最小值"]), "gsrc": d["收支來源"] or "", "alt": num(d["對照口徑空間"]), "altlbl": d["對照口徑"] or "",
            "pe28": d["現價 ÷ 模型 FY2028 EPS（A／B／C）"] or "", "note": d["註記"] or "",
            "c1src": d["中心候選 1 來源"] or "", "c2src": d["中心候選 2 來源"] or "",
            "b1src": d["悲觀情境候選 1 來源"] or "", "b2src": d["悲觀情境候選 2 來源"] or "",
            "tier6": d.get("上檔空間級距") or "", "recsrc": d.get("投資建議與上檔空間依據") or "",
        })

    nodes = [{"seg": d["次產業"], "order": d["次產業內優先序"], "name": d["公司"], "tk": str(d["代號"] or ""),
              "listing": d["掛牌"], "role": d["角色"] or "", "imp": d["供應鏈重要性（1–3）"],
              "pos": d["市場地位"] or "", "tag": d["標記"] or "", "status": d["研究狀態"] or ""}
             for d in rows(wb["地圖節點"])]

    params = {r[0]: r[1] for r in wb["參數"].iter_rows(min_row=2, values_only=True) if r[0]}

    data = {
        "meta": {"version": os.path.basename(src).split("_")[-1].replace(".xlsx", ""),
                 "book": os.path.basename(src), "tk_file": snap["file"], "tk_version": snap["version"],
                 "tk_commit": snap["commit"][:7], "cm_commit": cm, "built": "2026-10-26",
                 "price_date": "半導體：各模型現價（2026-10-22／23）；雲端與模型商：2026-10-23 收盤",
                 "pages": "https://yinchenchang.github.io/company-models/",
                 "private_models": {k: os.path.basename(glob.glob(os.path.join(ROOT, k, "dist", "*.html"))[0])
                                    for k in ("openai", "anthropic") if glob.glob(os.path.join(ROOT, k, "dist", "*.html"))}},
        "pergw": pergw, "segments": segments, "position": position, "valuation": valuation,
        "nodes": nodes, "params": params,
        "profiles": load_all(os.path.join(HERE, "profiles")),
        "models": model_files(valuation),
    }
    miss = {k: v["missing"] for k, v in data["profiles"].items() if v["missing"]}
    if miss:
        print("警告：公司檔案缺少段落", miss)
    nop = [v["tk"] for v in valuation if v["tk"] not in data["profiles"]]
    if nop:
        print("警告：尚無公司檔案", nop)
    tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
    # 造訪計數器網址（Google Apps Script 網頁應用程式；counter/endpoint.txt，空白＝不計數）
    ep = os.path.join(HERE, "counter", "endpoint.txt")
    cnt = open(ep, encoding="utf-8").read().strip() if os.path.exists(ep) else ""
    if cnt and not re.match(r"^https://script\.google\.com/macros/s/[\w-]+/exec$", cnt):
        raise SystemExit(f"counter/endpoint.txt 格式不對：{cnt}")
    html = tpl.replace("/*__DATA__*/null", json.dumps(data, ensure_ascii=False)).replace('/*__COUNTER__*/""', json.dumps(cnt))
    out = os.path.join(HERE, "dist", "index.html")
    with open(out, "w", encoding="utf-8") as f:
        f.write(html)
    with open(os.path.join(HERE, "dist", "data.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print("寫出", out, f"（{len(valuation)} 檔估值、{len(position)} 檔地位、{len(nodes)} 個地圖節點）")


def model_files(valuation):
    """各公司模型成品 HTML（與抽數用的 xlsx 同名）：網站「我們的模型」分頁嵌入。
    GitHub Pages 上的路徑為 models/<資料夾>/<檔名>（publish_pages.sh 複製）；本機開 dist/index.html 時改用 repo 相對路徑。"""
    ROOT = os.path.dirname(os.path.dirname(HERE))
    ext = json.load(open(os.path.join(HERE, "data", "models_extract.json"), encoding="utf-8"))
    out = {}
    for v in valuation:
        m = (ext.get(v["tk"]) or {}).get("model")
        if not m:
            continue
        h = os.path.splitext(m)[0] + ".html"
        if not os.path.exists(os.path.join(ROOT, h)):
            print("警告：找不到模型 HTML", h)
            continue
        folder, _, fn = h.split("/", 2)
        stem = fn[:-5]
        mm = re.match(r"^(\d{8})_(.+?)_(v[\d._]+)$", stem)
        date, title, ver = (mm.group(1), mm.group(2), mm.group(3).replace("_", ".")) if mm else ("", stem, "")
        out[v["tk"]] = {"folder": folder, "file": fn, "title": title,
                        "date": f"{date[:4]}-{date[4:6]}-{date[6:]}" if date else "", "ver": ver,
                        "xlsx": os.path.basename(m)}
    return out


if __name__ == "__main__":
    main()
