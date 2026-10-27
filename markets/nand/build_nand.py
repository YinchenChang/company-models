# NAND 市場層 v0.1（AI 半導體 Project）：CY2024–CY2030 位元需求（企業級 SSD、手機、PC、消費）、位元供給、平均售價與市場營收、各廠份額
# 用法：python3 markets/nand/build_nand.py raw.xlsx ；再以 LibreOffice headless 重算（錯誤格須為 0）
# 樣板：markets/memory（DRAM／HBM 市場層）。價格路徑、加速器顆數與情境顆數係數 SCEN_UNITS 一律讀 markets/memory/market.py（單一來源，不自訂行業路徑）。
# 情境 A／B／C 讀同字母：NAND 合約價季增（markets/memory NAND_S）、2028–29 加速器顆數係數（SCEN_UNITS）；CY2030 價格為收斂法（十年平均年變動）。
# 下游：sndk/、kioxia/ 讀本檔成品（dist）的 Mkt 頁季度營收與位元列；兩家份額合計＝本檔「Kioxia＋SanDisk（合資群）」列。
import sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
import market as MEM

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); FADE = PatternFill("solid", fgColor="E7E6E6")
SC = ("A", "B", "C")
SCN = {"A": "順利", "B": "平平", "C": "不如預期"}
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
def defname(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)

QS = [f"{y}Q{q}" for y in range(2025, 2031) for q in (1, 2, 3, 4)]       # 24 季：C..Z
QCOL = {q: openpyxl.utils.get_column_letter(3 + i) for i, q in enumerate(QS)}
YRS = list(range(2024, 2031))                                            # 年：C..I
YC = {y: chr(ord("C") + y - 2024) for y in YRS}
ACT = QS[:6]                                                              # 2025Q1–2026Q2 實績
VEN = [("ss", "三星"), ("sk", "SK 海力士＋Solidigm"), ("mu", "美光"), ("kx", "Kioxia"), ("sd", "SanDisk"), ("ot", "其他（長江存儲等）")]
TOP5 = ["ss", "sk", "mu", "kx", "sd"]

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — NAND 市場層輸入（藍字；基準 E 欄，低 D、高 F）。金額：十億美元（$B）；位元：EB（10^18 位元組）；顆數：百萬顆。每個輸入附標記與來源；找不到者寫「未找到」"
IN["A2"] = "標記：Verified／Interested-party（寫明是誰）／Analogy／Assumed（給低／高區間）／Derived（寫明換算）。價格路徑、加速器顆數、情境顆數係數取自 markets/memory/market.py（單一來源，本頁只抄錄並標示來源；改 markets/memory 後重建本檔）。SemiAnalysis 未使用"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
MEMSRC = "markets/memory/market.py"
TF0818 = "TrendForce 2026-08-18（2Q26 NAND 前五大營收排名）https://www.trendforce.com/presscenter/news/20260818-13186.html"

sec("N1. NAND 合約價季增（共用市場層 markets/memory；情境 A／B／C 同字母）")
for i, v in enumerate(MEM.NAND_COMMON):
    inp(f"nq_{i+1}", f"NAND 合約價季增 {MEM.QUARTERS[i+1]}", "%", None, v, None, "Interested-party", f"{MEMSRC} NAND_COMMON[{i}]", MEM.NAND_SRC if i == 0 else "同上")
for s in SC:
    for i, v in enumerate(MEM.NAND_S[s]):
        inp(f"nq{s}_{i+8}", f"NAND 合約價季增 {MEM.QUARTERS[i+8]}（{s}）", "%", None, v, None, "Interested-party" if s != "C" else "Assumed", f"{MEMSRC} NAND_S['{s}'][{i}]", MEM.NAND_S_SRC[s] if i == 0 else "同上")
inp("p10", "CY2030 NAND 平均售價年變動（收斂法：十年平均；三情境相同）", "%", -0.20, -0.12, -0.05, "Analogy",
    "上游 markets/memory 季價只到 2029Q4。收斂法：CY2030 起價格回到長期趨勢＝過去十年 NAND 每 GB 售價的年均變動；可驗證的 2016–2025 年 $/GB 序列未找到，取每位元成本年降 15–20% 減去循環中段的溢價（Analogy：製程每代密度提升、價格長期跟隨成本）",
    "2030 各季季增＝(1＋p10)^(1/4)−1，自 2029Q4 起算；低 −20%、高 −5%（Sens 列）")

sec("N2. 加速器顆數（markets/memory UNITS 合計；2028–29 情境係數 SCEN_UNITS，同字母）")
for i, y in enumerate((26, 27, 28)):
    tot = sum(u[2][i] for u in MEM.UNITS)
    inp(f"u_{y}", f"加速器顆數合計 20{y}（情境 A 基準）", "百萬顆", None, round(tot, 6), None, "Derived", f"{MEMSRC} UNITS 各列 20{y} 合計（NVIDIA、AMD、TPU、Trainium、其他 ASIC）", "與 HBM 需求、台積電 CoWoS、Advantest 同一套顆數" if i == 0 else "同上")
for s in SC:
    for n, y in enumerate((28, 29)):
        inp(f"um{y}_{s}", f"情境 {s}：20{y} 顆數 ÷ 基準", "x", None, MEM.SCEN_UNITS[s][n], None, "Derived", f"{MEMSRC} SCEN_UNITS['{s}'][{n}]", "A＝1.0；B、C 為 Assumed（markets/memory）；CY2030 沿用 2029 係數" if (s == "A" and y == 28) else "同上")
inp("g_u29", "2029 加速器顆數年增（情境 A 基準；B、C 再乘 um29）", "%", 0.0, MEM.G_U29, 0.25, "Derived", "Siliconomics SIF_UnitsGrowth29：平台別 2029 成長率（M5 節）加權，與台積電、網通模型同一路徑（v0.2 起；原 Assumed 10%）；低／高沿用原區間供敏感度", "三情境同一基準，情境差異只來自 um29（不自訂行業路徑）")
inp("g_u30", "2030 加速器顆數年增（情境 A 基準；B、C 再乘 um29）", "%", -0.05, 0.05, 0.15, "Assumed", "上游只到 2029；收斂到低個位數成長", "同上")
for y, (lo, b, hi) in zip((26, 27, 28, 29, 30), ((24, 30, 36), (27, 34, 41), (32, 40, 48), (37, 46, 55), (42, 52, 62))):
    inp(f"tb_{y}", f"每顆加速器帶動的企業級 SSD 容量 20{y}", "TB／顆", lo, b, hi, "Assumed",
        "AI 推論的 KV 快取卸載、RAG 向量庫、檢查點與資料湖（伺服器內建＋外接儲存）。2026 值使 2026 企業級 SSD 約占位元需求四成（TrendForce 2026-07：伺服器約 40%；Counterpoint 2Q26：企業級 SSD 占出貨位元 48%）" if y == 26 else "每年約 +13–18%：單顆加速器記憶體與脈絡長度增加（Assumed）",
        "低／高 ±20%（Sens 列）；三情境相同（情境差異來自顆數）")

sec("N3. 位元需求（2025 基期與各應用年增；非 AI 應用三情境相同）")
inp("bits25", "2025 NAND 位元出貨（全產業）", "EB", 850, 1000, 1150, "Assumed", "全產業 EB 的公開數字未找到（TrendForce、Yole 付費）；以 2025 營收約 $72B ÷ 合約均價約 $0.07／GB 粗估約 1 ZB", "只影響 EB 與 $/GB 的絕對值，以及 AI 需求（顆數 × 容量，絕對量）與其他應用（比例 × 基期）的相對大小；營收以 2Q26 實績為錨，不受基期水準影響")
inp("g25", "2025 位元出貨年增（推 2024 位元）", "%", 0.08, 0.12, 0.16, "Assumed", "2025 上半庫存去化、下半價漲；全年位元成長公開數字未找到", "只用於 CY2024 欄")
inp("sh_essd25", "2025 企業級 SSD 占位元需求", "%", 0.26, 0.30, 0.34, "Interested-party", "Counterpoint 2Q25 企業級 SSD 占 NAND 出貨位元 26%、2Q26 48%（capitalandcompute 2026-08 轉載 Counterpoint Q2 2026 Memory & Storage Tracker）", "全年平均取 30%（下半年 AI 伺服器拉貨，Derived／Assumed）")
inp("ai25", "2025 企業級 SSD 中 AI 伺服器的比例", "%", 0.45, 0.55, 0.65, "Assumed", "未找到；AI 伺服器（加速器）帶動的企業級 SSD 約過半", "2026 起 AI 部分改由顆數 × 每顆容量推導")
inp("sh_mob25", "2025 手機占位元需求", "%", 0.28, 0.32, 0.36, "Assumed", "TrendForce 2026-07：手機與筆電合計約 40%（2026，企業級 SSD 擴大後）；2025 手機約三成", "")
inp("sh_pc25", "2025 PC（client SSD）占位元需求", "%", 0.20, 0.23, 0.26, "Assumed", "同上", "消費性（記憶卡、USB、遊戲機、車用等）＝1 − 其餘")
inp("g_nai", "非 AI 企業級 SSD 位元年增（一般伺服器、儲存陣列取代 HDD）", "%", 0.05, 0.10, 0.15, "Assumed", "QLC 取代近線 HDD（TrendForce 2025-09：HDD 交期逾 52 週）", "三情境相同")
for k, nm, vals, src in (("mob", "手機", ((-0.10, -0.05, 0.0), (0.0, 0.03, 0.06), (0.05, 0.08, 0.11)), "TrendForce 2026-08-18：手機、PC 因 BOM 成本上升需求轉弱；2028 起每機容量回升"),
                         ("pc", "PC", ((-0.10, -0.05, 0.0), (-0.03, 0.0, 0.03), (0.03, 0.06, 0.09)), "同上"),
                         ("con", "消費性與其他", ((-0.15, -0.10, -0.05), (-0.03, 0.0, 0.03), (0.02, 0.05, 0.08)), "微軟、任天堂 2026 年遊戲機漲價（儲存成本上升 2.5 倍）")):
    for (lo, b, hi), lab in zip(vals, ("26", "27", "28")):
        inp(f"g_{k}_{lab}", f"{nm}位元年增 20{lab}" + ("（2029–30 同此值）" if lab == "28" else ""), "%", lo, b, hi, "Assumed", src if lab == "26" else "同上", "三情境相同（非 AI 應用不隨加速器情境）")

sec("N4. 位元供給（全產業；三情境相同：產能與製程轉換 2027 前已訂定）")
for y, (lo, b, hi, tag, src) in zip((26, 27, 28, 29, 30), ((0.18, 0.21, 0.24, "Interested-party", "美光 FQ4 FY26（2026-09-30）：產業 NAND 位元出貨 CY2026 低 20% 段、CY2027 與 CY2028 約中 20% 段，供給受限至 2028（Guru3D 2026-10-01 轉載）"),
                                                   (0.20, 0.25, 0.28, "Interested-party", "同上"), (0.20, 0.25, 0.28, "Interested-party", "同上"),
                                                   (0.15, 0.22, 0.28, "Assumed", "上游無 2029；Kioxia 北上 K2（BiCS10，332 層）、三星 V10（430 層以上）2027 起量產"),
                                                   (0.12, 0.20, 0.26, "Assumed", "同上"))):
    inp(f"gs_{y}", f"位元供給年增 20{y}", "%", lo, b, hi, tag, src, "三星 2026 晶圓投片約 −4.5%、SK 海力士 −10%（投資轉向 HBM；borecraft 2026-05-26 轉載 DigiTimes、Blocks & Files），位元成長來自層數" if y == 26 else "同上")

sec("N5. 季度實績（TrendForce 前五大 NAND 營收；$B）與前五大占全市場比例")
RV = {"ss": (4.20, 5.20, 6.00, 6.60, 13.51, 23.06), "sk": (2.19, 3.34, 3.53, 5.21, 7.53, 14.27), "mu": (2.03, 2.10, 2.42, 3.03, 5.95, 11.85),
      "kx": (1.92, 2.14, 2.84, 3.31, 5.96, 10.72), "sd": (1.69, 1.90, 2.31, 3.03, 5.95, 8.97)}
RVSRC = {0: ("Derived", "2Q25 營收 ÷（1＋2Q25 季增）（TrendForce 2025-08-28：三星 +23.8%、SK +52.5%、Kioxia +11.4%、美光 +3.7%、SanDisk +12.2%）"),
         1: ("Interested-party", "TrendForce 2025-08-28（2Q25 前五大 $14.67B）https://www.trendforce.com/presscenter/news/20250828-12688.html"),
         2: ("Interested-party", "TrendForce 2025-12-03（3Q25 前五大約 $17.1B）https://www.trendforce.com/presscenter/news/20251203-12813.html"),
         3: ("Interested-party", "TrendForce（4Q25 前五大 $21.17B，+23.8%；I-Connect007 轉載）"),
         4: ("Derived", "2Q26 營收 ÷（1＋2Q26 季增）（TrendForce 2026-08-18：三星 +70.7%、SK +89.5%、美光 +99.2%、Kioxia +79.9%、SanDisk +50.7%）"),
         5: ("Interested-party", TF0818)}
for k, nm in VEN[:5]:
    for i, q in enumerate(ACT):
        tag, src = RVSRC[i]
        inp(f"rv_{k}_{i}", f"{nm} NAND 營收 {q}", "$B", None, RV[k][i], None, tag, src,
            ("SanDisk 公司營收：FQ1 FY26 2.308、FQ2 3.025、FQ3 5.950、FQ4 8.965（8-K），與 TrendForce 一致" if (k == "sd" and i == 5) else ""))
for i, (q, v, lo, hi, tag, src) in enumerate(((ACT[0], 0.93, 0.90, 0.95, "Assumed", "未找到；取 2Q25 值附近"),
                                              (ACT[1], "=(rv_ss_1+rv_sk_1+rv_mu_1+rv_kx_1+rv_sd_1)/(rv_ss_1/0.329)", None, None, "Derived", "全市場＝三星營收 ÷ 三星市占 32.9%（TrendForce 2025-08-28）；前五大 ÷ 全市場"),
                                              (ACT[2], 0.91, 0.88, 0.93, "Assumed", "未找到；2Q25 與 4Q25 之間"),
                                              (ACT[3], "=(rv_ss_3+rv_sk_3+rv_mu_3+rv_kx_3+rv_sd_3)/(rv_ss_3/0.28)", None, None, "Derived", "三星市占 28%（TrendForce 4Q25）"),
                                              (ACT[4], 0.885, 0.86, 0.90, "Assumed", "未找到；4Q25 與 2Q26 之間"),
                                              (ACT[5], "=(rv_ss_5+rv_sk_5+rv_mu_5+rv_kx_5+rv_sd_5)/(rv_ss_5/0.293)", None, None, "Derived", "三星市占 29.3%、Kioxia 13.6%（TrendForce 2026-08-18；兩者推得的全市場一致）"))):
    inp(f"t5_{i}", f"前五大占全市場營收 {q}", "%", lo, v, hi, tag, src, "其他＝長江存儲、英特爾殘餘、旺宏、華邦等")
inp("rev24", "2024 NAND 市場營收", "$B", None, 67.4, None, "Interested-party", "TrendForce（TelecomLead 轉載）：2024 NAND 營收 $67.4B（+77%）", "第二來源：未找到；只用於 CY2024 欄")

sec("N6. 各廠份額（2026Q3 起；Kioxia、SanDisk 由公司 2026Q3 指引推得，其餘沿用 1H26 實績，其他＝餘額）")
inp("sd_g", "SanDisk FQ1 FY27（2026-07～09）營收指引中值", "$B", 10.3, 10.55, 10.8, "Interested-party", "SanDisk FQ4 FY26 新聞稿（2026-08-05）：$10.3–10.8B https://www.sec.gov/Archives/edgar/data/0002023554/000162828026053346/sndkq4-26ex991xpressrelease.htm", "公司指引（立場：公司）")
inp("sd_k", "SanDisk 公司營收 ÷ TrendForce 口徑（2Q26）", "x", None, "=8.965/rv_sd_5", None, "Derived", "FQ4 FY26 營收 8.965 ÷ TrendForce 2Q26 8.97", "")
inp("kx_g", "Kioxia FY2026 第二季（2026-07～09）營收指引", "¥B", None, 2390.0, None, "Interested-party", "Kioxia FY2026 第一季財報（2026-07-31，SGX 轉載 Earnings Release FY2026 1Q）：第二季營收 ¥2,390.0B、營業利益 ¥1,890.0B，假設 1 美元＝162 日圓", "公司指引（立場：公司）")
inp("kx_fxg", "Kioxia 第二季指引匯率", "¥／$", None, 162, None, "Interested-party", "同上", "")
inp("kx_k", "Kioxia 公司營收 ÷ TrendForce 口徑（2Q26）", "x", None, "=1767.1/160/rv_kx_5", None, "Derived", "第一季營收 ¥1,767.1B ÷ 平均匯率 160 ÷ TrendForce 2Q26 10.72（公司含其他營收 ¥66.7B）", "")
for k, (lo, b, hi), src in (("ss", (-1.0, -0.5, 0.0), "三星：2026 晶圓投片減少、投資轉向 HBM（Blocks & Files 2026-01-20 轉載）；2Q26 市占小降（TrendForce）"),
                            ("sk", (-0.5, 0.0, 0.5), "SK 海力士＋Solidigm：投片 −10%，但 QLC 企業級 SSD 接單強（TrendForce 2026-08-18）"),
                            ("mu", (-0.3, 0.2, 0.6), "美光：2Q26 季增最高（+99.2%）；位元成長低於產業（美光 FQ4 FY26）")):
    inp(f"dr_{k}", f"{dict(VEN)[k]} 營收份額年變動（2027 起，三情境相同）", "pt／年", lo, b, hi, "Assumed", src, "")
for s, (lo, b, hi), src in zip(SC, ((-0.3, 0.0, 0.3), (-0.6, -0.3, 0.0), (-1.0, -0.6, -0.2)),
                               ("A：北上 K2 與 BiCS10 如期量產（Kioxia FY26–28 年均資本支出 ¥470B，TrendForce 2026-06-03）", "B：長江存儲（2Q26 位元市占 14%，Counterpoint）持續擴產", "C：AI 需求轉弱時中國廠以價換量")):
    inp(f"dr_jv_{s}", f"情境 {s}：Kioxia＋SanDisk 合資群營收份額年變動（2027 起）", "pt／年", lo, b, hi, "Assumed", src,
        "Flash Ventures（四日市、北上）由 Kioxia 50.1%／SanDisk 49.9% 持有（SanDisk FQ4 FY26 簡報），兩家位元成長同步；方向：A 最有利、C 最不利；其他（長江存儲等）吸收差額" if s == "A" else "同上")

IN_ROW = {}; r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}"); IN_ROW[key] = r
    if lo is not None: defname(key + "_lo", f"In!$D${r}")
    if hi is not None: defname(key + "_hi", f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 60, 10, 9, 11, 9, 18, 90, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# ---------------- Path：季價指數（2025Q1＝1）----------------
P = wb.create_sheet("Path")
P["A1"] = "Path — NAND 合約價指數（2025Q1＝1）。2025Q2–2029Q4 季增取自 markets/memory（In 頁 N1）；2030 各季＝(1＋p10)^(1/4)（收斂法，灰底）"
hdr(P, 3, ["情境", "說明"] + QS)
PR = {}
for i, s in enumerate(SC):
    rr = 4 + i; PR[s] = rr
    P.cell(rr, 1, s); P.cell(rr, 2, f"情境 {s}（{SCN[s]}）")
    P.cell(rr, 3, 1)
    for j in range(1, 24):
        c = openpyxl.utils.get_column_letter(3 + j); pc = openpyxl.utils.get_column_letter(2 + j)
        if j <= 7: g = f"nq_{j}"
        elif j <= 19: g = f"nq{s}_{j}"
        else: g = "((1+p10)^(1/4)-1)"
        P[f"{c}{rr}"] = f"={pc}{rr}*(1+{g})"; P[f"{c}{rr}"].number_format = "0.000"
        if j >= 20: P[f"{c}{rr}"].fill = FADE
P.cell(8, 1, "年均指數").font = BOLD
hdr(P, 9, ["情境", "說明"] + [f"CY{y}" for y in range(2025, 2031)] + [f"CY{y}÷CY{y-1}" for y in range(2026, 2031)])
for i, s in enumerate(SC):
    rr = 10 + i; P.cell(rr, 1, s); P.cell(rr, 2, "四季平均")
    for j, y in enumerate(range(2025, 2031)):
        a, b = QCOL[f"{y}Q1"], QCOL[f"{y}Q4"]
        P.cell(rr, 3 + j, f"=AVERAGE({a}{PR[s]}:{b}{PR[s]})").number_format = "0.000"
    for j in range(5):
        c1 = openpyxl.utils.get_column_letter(4 + j); c0 = openpyxl.utils.get_column_letter(3 + j)
        P.cell(rr, 9 + j, f"={c1}{rr}/{c0}{rr}-1").number_format = "0.0%"
P.column_dimensions["B"].width = 18

# ---------------- 引擎：一個情境的完整推導（Mkt 與 Sens 共用） ----------------
DEFP = dict(tb="", gs="", u29="g_u29", u30="g_u30", p10="p10", bits="bits25", essd="sh_essd25", ai="ai25", nai="g_nai", px="1", jv=None)
def engine(ws, r0, title, s, **over):
    p = dict(DEFP); p.update(over)
    jv = p["jv"] or f"dr_jv_{s}"
    R = {}; rr = [r0]
    ws.cell(r0, 1, title).font = BOLD; rr[0] += 1
    def row(key, label, unit, fs, fmt="#,##0.0", note="", ncol=None):
        r_ = rr[0]; ws.cell(r_, 1, label); ws.cell(r_, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r_}"] = f; ws[f"{col}{r_}"].number_format = fmt
        if note: ws.cell(r_, ncol or 28, note)
        R[key] = r_; rr[0] += 1
    def sub(t):
        ws.cell(rr[0], 1, t).font = BOLD; rr[0] += 1
    hdr(ws, rr[0], ["A. 年度（日曆年）", "單位"] + [f"CY{y}" for y in YRS] + [""] * 18 + ["說明"]); rr[0] += 1
    Y6 = range(2026, 2031)
    um = {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}", 2030: f"um29_{s}"}
    base = {2026: "u_26", 2027: "u_27", 2028: "u_28", 2029: f"u_28*(1+{p['u29']})", 2030: f"u_28*(1+{p['u29']})*(1+{p['u30']})"}
    row("u", "加速器顆數（同字母情境係數）", "百萬顆", {YC[y]: f"=({base[y]})*{um[y]}" for y in Y6}, note="markets/memory UNITS × SCEN_UNITS；2029–30 基準以 g_u29、g_u30 外推，CY2030 沿用 um29")
    row("tb", "每顆企業級 SSD 容量", "TB", {YC[y]: f"=tb_{y % 100}{p['tb']}" for y in Y6}, fmt="0.0")
    row("dai", "需求：AI 企業級 SSD", "EB", {YC[2025]: f"={p['bits']}*{p['essd']}*{p['ai']}", **{YC[y]: f"={YC[y]}{R['u']}*{YC[y]}{R['tb']}" for y in Y6}}, note="顆數 × 每顆容量（百萬顆 × TB＝EB）；2025＝基期 × 占比")
    row("dne", "需求：非 AI 企業級 SSD", "EB", {YC[2025]: f"={p['bits']}*{p['essd']}*(1-{p['ai']})", **{YC[y]: f"={YC[y-1]}{rr[0]}*(1+{p['nai']})" for y in Y6}})
    gg = lambda k, y: f"g_{k}_{min(y, 2028) % 100}"
    row("dmo", "需求：手機", "EB", {YC[2025]: f"={p['bits']}*sh_mob25", **{YC[y]: f"={YC[y-1]}{rr[0]}*(1+{gg('mob', y)})" for y in Y6}})
    row("dpc", "需求：PC（client SSD）", "EB", {YC[2025]: f"={p['bits']}*sh_pc25", **{YC[y]: f"={YC[y-1]}{rr[0]}*(1+{gg('pc', y)})" for y in Y6}})
    row("dco", "需求：消費性與其他", "EB", {YC[2025]: f"={p['bits']}*(1-{p['essd']}-sh_mob25-sh_pc25)", **{YC[y]: f"={YC[y-1]}{rr[0]}*(1+{gg('con', y)})" for y in Y6}})
    row("dem", "位元需求合計", "EB", {YC[y]: f"=SUM({YC[y]}{R['dai']}:{YC[y]}{R['dco']})" for y in range(2025, 2031)})
    row("ess", "企業級 SSD 占位元需求", "%", {YC[y]: f"=({YC[y]}{R['dai']}+{YC[y]}{R['dne']})/{YC[y]}{R['dem']}" for y in range(2025, 2031)}, fmt="0.0%")
    row("sup", "位元供給（三情境相同）", "EB", {YC[2025]: f"={p['bits']}", **{YC[y]: f"={YC[y-1]}{rr[0]}*(1+gs_{y % 100}{p['gs']})" for y in Y6}}, note="2025＝出貨實績（基期）；供給不隨需求情境變（產能 2027 前已訂定），需求不足的結果是閒置產能與降價")
    row("gap", "需求 ÷ 供給 − 1（>0＝短缺）", "%", {YC[y]: f"={YC[y]}{R['dem']}/{YC[y]}{R['sup']}-1" for y in Y6}, fmt="0.0%")
    row("ship", "位元出貨＝MIN(需求, 供給)", "EB", {YC[2024]: f"={p['bits']}/(1+g25)", YC[2025]: f"={p['bits']}", **{YC[y]: f"=MIN({YC[y]}{R['dem']},{YC[y]}{R['sup']})" for y in Y6}})
    row("bg", "位元出貨年增", "%", {YC[y]: f"={YC[y]}{R['ship']}/{YC[y-1]}{R['ship']}-1" for y in range(2025, 2031)}, fmt="0.0%")
    yr_rows = rr[0]; rr[0] += 11   # 年度營收等列在季度區塊之後回填
    # B. 季度
    hdr(ws, rr[0], ["B. 季度", "單位"] + QS + ["", "說明"]); rr[0] += 1
    qy = lambda q: int(q[:4]); qk = lambda q: int(q[-1])
    row("pq", "合約價指數（同字母路徑）", "2025Q1＝1", {QCOL[q]: (f"=Path!{QCOL[q]}{PR[s]}" if qy(q) < 2030 else f"=Path!{QCOL['2029Q4']}{PR[s]}*(1+{p['p10']})^({qk(q)}/4)") + (f"*{p['px']}" if qy(q) >= 2027 else "") for q in QS}, fmt="0.000", note="Path 頁；Sens 可對 2027 起的季價乘價格係數")
    def bq(q):
        y = qy(q); k = qk(q); g = f"({YC[y]}{R['ship']}/{YC[y-1]}{R['ship']})"
        return f"={YC[y]}{R['ship']}/4*{g}^(({k}-2.5)/4)/(({g}^(-0.375)+{g}^(-0.125)+{g}^(0.125)+{g}^(0.375))/4)"
    row("bq", "季度位元出貨（年內依年增平滑；四季合計＝年度）", "EB", {QCOL[q]: bq(q) for q in QS}, note="季度＝年度 ÷ 4 ×（1＋年增）^((季−2.5)/4)，再正規化使四季合計等於年度")
    a6 = QCOL[ACT[5]]
    def mq(q):
        if q in ACT:
            i = ACT.index(q); return "=(" + "+".join(f"rv_{k}_{i}" for k in TOP5) + f")/t5_{i}"
        return f"=${a6}{rr[0]}*{QCOL[q]}{R['pq']}/${a6}{R['pq']}*{QCOL[q]}{R['bq']}/${a6}{R['bq']}"
    row("mq", "NAND 市場營收（2026Q2 前實績；之後＝2Q26 實績 × 價格指數比 × 位元比）", "$B", {QCOL[q]: mq(q) for q in QS}, note="實績＝TrendForce 前五大 ÷ 前五大占比；以 2Q26 為錨")
    q3 = QCOL["2026Q3"]
    # 份額（季）
    for k, nm in VEN:
        def shq(q):
            if q in ACT:
                i = ACT.index(q)
                return f"=rv_{k}_{i}/{QCOL[q]}{R['mq']}" if k != "ot" else f"=1-t5_{i}"
            y = qy(q)
            if k == "kx": b0 = f"kx_g/kx_fxg/kx_k/${q3}{R['mq']}"
            elif k == "sd": b0 = f"sd_g/sd_k/${q3}{R['mq']}"
            elif k == "ot": return f"=1-" + "-".join(f"{QCOL[q]}{R['sh_' + kk]}" for kk in TOP5)
            else: b0 = f"(rv_{k}_4+rv_{k}_5)/(${QCOL['2026Q1']}{R['mq']}+${a6}{R['mq']})"
            if y == 2026: return "=" + b0
            d = {"ss": "dr_ss", "sk": "dr_sk", "mu": "dr_mu"}.get(k)
            if d: return f"={b0}+{d}/100*{y - 2026}"
            frac = f"(kx_g/kx_fxg/kx_k)/(kx_g/kx_fxg/kx_k+sd_g/sd_k)" if k == "kx" else f"(sd_g/sd_k)/(kx_g/kx_fxg/kx_k+sd_g/sd_k)"
            return f"={b0}+{jv}/100*{y - 2026}*{frac}"
        R["sh_" + k] = rr[0]
        row("sh_" + k, f"份額：{nm}", "%", {QCOL[q]: shq(q) for q in QS}, fmt="0.0%",
            note={"kx": "2026Q3＝公司第二季指引 ÷ 匯率 ÷ 公司口徑係數 ÷ 市場；2027 起＋合資群年變動 × Kioxia 比重", "sd": "2026Q3＝公司 FQ1 FY27 指引 ÷ 口徑係數 ÷ 市場；2027 起＋合資群年變動 × SanDisk 比重",
                  "ot": "餘額（長江存儲等）"}.get(k, "2026Q3–Q4＝1H26 實績份額；2027 起每年＋年變動"))
    for k, nm in VEN:
        row("rv_" + k, f"{nm} 營收（情境 {s}）", "$B", {QCOL[q]: f"={QCOL[q]}{R['mq']}*{QCOL[q]}{R['sh_' + k]}" for q in QS}, fmt="#,##0.00")
    row("jvq", f"Kioxia＋SanDisk（合資群）份額（情境 {s}）", "%", {QCOL[q]: f"={QCOL[q]}{R['sh_kx']}+{QCOL[q]}{R['sh_sd']}" for q in QS}, fmt="0.0%")
    for k in ("kx", "sd"):
        row("bt_" + k, f"{dict(VEN)[k]} 位元（情境 {s}；季度出貨 × 營收份額）", "EB", {QCOL[q]: f"={QCOL[q]}{R['bq']}*{QCOL[q]}{R['sh_' + k]}" for q in QS}, fmt="#,##0.00",
            note="假設該廠平均售價＝市場平均（份額以營收計）；公司模型以此列的變動率推單位成本" if k == "kx" else "")
    row("chk", "檢查：份額合計 − 1；實績季營收重現（應為 0）", "—", {"C": "=" + "+".join(f"ABS(SUM({QCOL[q]}{R['sh_ss']}:{QCOL[q]}{R['sh_ot']})-1)" for q in (QS[0], QS[6], QS[-1])) + f"+ABS({QCOL[ACT[5]]}{R['rv_sd']}-rv_sd_5)"}, fmt="0.000000")
    end = rr[0]
    # 回填年度列
    rr[0] = yr_rows
    def ysum(rowkey, y): return f"=SUM({QCOL[f'{y}Q1']}{R[rowkey]}:{QCOL[f'{y}Q4']}{R[rowkey]})"
    row("rev", "NAND 市場營收（四季合計）", "$B", {YC[2024]: "=rev24", **{YC[y]: ysum("mq", y) for y in range(2025, 2031)}}, note="CY2024 為 TrendForce 年度值")
    row("rg", "市場營收年增", "%", {YC[y]: f"={YC[y]}{R['rev']}/{YC[y-1]}{R['rev']}-1" for y in range(2025, 2031)}, fmt="0.0%")
    row("asp", "平均售價（營收 ÷ 位元出貨）", "$/GB", {YC[y]: f"={YC[y]}{R['rev']}/{YC[y]}{R['ship']}" for y in YRS}, fmt="0.000")
    row("aspg", "平均售價年變動", "%", {YC[y]: f"={YC[y]}{R['asp']}/{YC[y-1]}{R['asp']}-1" for y in range(2025, 2031)}, fmt="0.0%", note="CY2030 為收斂法（p10）＋位元平滑效果")
    for k, nm in VEN:
        row("ya_" + k, f"年度份額：{nm}", "%", {YC[y]: f"=SUM({QCOL[f'{y}Q1']}{R['rv_' + k]}:{QCOL[f'{y}Q4']}{R['rv_' + k]})/{YC[y]}{R['rev']}" for y in range(2025, 2031)}, fmt="0.0%")
    rr[0] = end
    R["end"] = end
    return R

M = wb.create_sheet("Mkt")
M["A1"] = "Mkt — NAND 市場（三情境各一個完整推導區塊；A 區年度、B 區季度）。需求＝AI 企業級 SSD（加速器顆數 × 每顆容量）＋非 AI；出貨＝MIN(需求, 供給)；營收＝2Q26 實績 × 價格指數比 × 位元比；份額 × 市場＝各廠營收"
M["A2"] = "下游讀數列：「Kioxia 營收（情境 s）」「SanDisk 營收（情境 s）」「Kioxia 位元（…）」「SanDisk 位元（…）」（C..Z＝2025Q1..2030Q4）。灰底＝2030 收斂法"
MB = {}; r0 = 4
for s in SC:
    MB[s] = engine(M, r0, f"情境 {s}（{SCN[s]}）：價格路徑 {s}（{MEM.SCEN_DESC[s]}）；2028–29 顆數 × {MEM.SCEN_UNITS[s][0]}／{MEM.SCEN_UNITS[s][1]}；合資群份額年變動 dr_jv_{s}", s)
    r0 = MB[s]["end"] + 2
    for rk in ("pq", "mq"):
        for q in QS[20:]: M[f"{QCOL[q]}{MB[s][rk]}"].fill = FADE
M.column_dimensions["A"].width = 62; M.column_dimensions["B"].width = 10
for q in QS: M.column_dimensions[QCOL[q]].width = 9
M.column_dimensions["AB"].width = 70
M.freeze_panes = "C4"

# ---------------- Ann：年度摘要（三情境） ----------------
A = wb.create_sheet("Ann", 1)
A["A1"] = "Ann — NAND 市場年度摘要（三情境；引用 Mkt 頁）。CY2024–2025 為實績（位元為 Assumed 基期）；CY2026 含 1H26 實績；CY2030 價格為收斂法（灰底）"
hdr(A, 3, ["情境／項目", "單位"] + [f"CY{y}" for y in YRS] + ["說明"])
AR = {}; ar = 4
ITEMS = [("u", "加速器顆數", "百萬顆", "#,##0.0"), ("dai", "  需求：AI 企業級 SSD", "EB", "#,##0"), ("dne", "  需求：非 AI 企業級 SSD", "EB", "#,##0"), ("dmo", "  需求：手機", "EB", "#,##0"),
         ("dpc", "  需求：PC", "EB", "#,##0"), ("dco", "  需求：消費性與其他", "EB", "#,##0"), ("dem", "位元需求合計", "EB", "#,##0"), ("ess", "企業級 SSD 占需求", "%", "0.0%"),
         ("sup", "位元供給", "EB", "#,##0"), ("gap", "需求 ÷ 供給 − 1", "%", "0.0%"), ("ship", "位元出貨", "EB", "#,##0"), ("bg", "位元出貨年增", "%", "0.0%"),
         ("asp", "平均售價", "$/GB", "0.000"), ("aspg", "平均售價年變動", "%", "0.0%"), ("rev", "NAND 市場營收", "$B", "#,##0.0"), ("rg", "市場營收年增", "%", "0.0%")] + \
        [("ya_" + k, "份額：" + nm, "%", "0.0%") for k, nm in VEN]
for s in SC:
    A.cell(ar, 1, f"情境 {s}（{SCN[s]}）").font = BOLD; ar += 1
    for key, lab, unit, fm in ITEMS:
        A.cell(ar, 1, lab); A.cell(ar, 2, unit)
        for y in YRS:
            src = M[f"{YC[y]}{MB[s][key]}"].value
            if src is None: continue
            A[f"{YC[y]}{ar}"] = f"=Mkt!{YC[y]}{MB[s][key]}"; A[f"{YC[y]}{ar}"].number_format = fm
            if y == 2030 and key in ("asp", "aspg", "rev", "rg"): A[f"{YC[y]}{ar}"].fill = FADE
        AR[(s, key)] = ar; ar += 1
    A.cell(ar, 1, "Kioxia＋SanDisk（合資群）份額"); A.cell(ar, 2, "%")
    for y in range(2025, 2031):
        A[f"{YC[y]}{ar}"] = f"={YC[y]}{AR[(s, 'ya_kx')]}+{YC[y]}{AR[(s, 'ya_sd')]}"; A[f"{YC[y]}{ar}"].number_format = "0.0%"
    AR[(s, "jv")] = ar; ar += 1
    A.cell(ar, 1, "Kioxia 營收（年）"); A.cell(ar, 2, "$B")
    A.cell(ar + 1, 1, "SanDisk 營收（年）"); A.cell(ar + 1, 2, "$B")
    for j, k in enumerate(("kx", "sd")):
        for y in range(2025, 2031):
            A[f"{YC[y]}{ar + j}"] = f"=SUM(Mkt!{QCOL[f'{y}Q1']}{MB[s]['rv_' + k]}:{QCOL[f'{y}Q4']}{MB[s]['rv_' + k]})"; A[f"{YC[y]}{ar + j}"].number_format = "#,##0.0"
    AR[(s, "kxr")] = ar; AR[(s, "sdr")] = ar + 1; ar += 3
A.cell(ar, 1, "核對：三情境份額合計 − 1、實績季重現（應為 0）"); A[f"C{ar}"] = "=" + "+".join(f"Mkt!C{MB[s]['chk']}" for s in SC); A[f"C{ar}"].number_format = "0.000000"; ACHK = ar
A.column_dimensions["A"].width = 40
for y in YRS: A.column_dimensions[YC[y]].width = 11
A.column_dimensions["J"].width = 50
A.cell(AR[("A", "rev")], 10, "CY2026–28 三情境營收與位元成長見 README")

# ---------------- Sup：各廠產能與層數（說明表） ----------------
SU = wb.create_sheet("Sup", 3)
SU["A1"] = "Sup — 主要廠位元供給說明（層數、產能動向；本頁為文字輸入，不進計算：位元供給以全產業年增 gs_ 計，見 In 頁 N4）"
hdr(SU, 3, ["廠商", "主力世代（2026）", "下一代", "2026 產能動向", "2Q26 位元市占", "標記", "來源"])
SUP = [("三星", "V9（286 層）", "V10（430 層以上；大規模投資不早於 1H26）", "晶圓投片約 −4.5%（投資轉向 HBM）", "25%", "Interested-party", "borecraft 2026-05-26（轉載 DigiTimes 2026-05-25、Blocks & Files 2026-01-20）；Counterpoint 2Q26 位元市占"),
       ("SK 海力士＋Solidigm", "321 層（V9 世代）", "300 層級新品 2027 年初全面量產", "晶圓投片約 −10%；Solidigm QLC 企業級 SSD 出貨季增約 40%", "22%", "Interested-party", "同上；TrendForce 2026-08-18"),
       ("Kioxia（合資群）", "BiCS8（218 層）", "BiCS10（332 層，北上 K2，FY2026 量產；密度高約 59%）", "Kioxia＋SanDisk 2026 資本支出 $4.5B（+41%）；Kioxia FY26–28 年均 ¥470B", "未找到", "Interested-party", "TrendForce 2025-11-13、2026-06-03；borecraft 2026-05-26"),
       ("SanDisk（合資群）", "BiCS8（218 層；與 Kioxia 共用晶圓廠）", "BiCS10", "位元成長相對保守（TrendForce 2Q26）；Flash Ventures 持股 49.9%", "未找到", "Interested-party", "TrendForce 2026-08-18；SanDisk FQ4 FY26 簡報"),
       ("美光", "G9（276 層）", "未找到", "位元成長低於產業（美光 FQ4 FY26）", "未找到", "Interested-party", "美光 FQ4 FY26 法說（Guru3D 2026-10-01）"),
       ("長江存儲", "Xtacking 4.0（267 層）", "未找到", "2Q26 位元市占 14%（第三），年增約 22%；三廠 2026 年底投產", "14%", "Interested-party", "Counterpoint 2Q26（hwbusters 轉載）；markets/equipment In 頁")]
for i, row_ in enumerate(SUP, 4):
    for j, v in enumerate(row_, 1): SU.cell(i, j, v)
SU.cell(11, 1, "各廠晶圓月產能（千片）：未找到可驗證的公開數字（TrendForce、Yole 付費）。本版不以各廠產能加總推供給，改用全產業位元年增（美光指引），各廠以營收份額分配。")
for c, w in zip("ABCDEFG", [22, 26, 40, 50, 12, 16, 70]): SU.column_dimensions[c].width = w

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、公開預測（只作事後對照，不回推參數）。差距＝模型 ÷ 對照 − 1（百分點列為相減）"
hdr(RC, 3, ["項目", "對照值", "模型 A", "模型 B", "模型 C", "差距 A", "對照來源", "說明"])
rcr = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False):
    global rcr
    RC.cell(rcr, 1, label); RC.cell(rcr, 2, ref).number_format = fmt
    for j, s in enumerate(SC): RC.cell(rcr, 3 + j, getter(s)).number_format = fmt
    RC.cell(rcr, 6, f"=C{rcr}-B{rcr}" if pt else f"=C{rcr}/B{rcr}-1").number_format = "0.0%"
    RC.cell(rcr, 7, src); RC.cell(rcr, 8, note); rcr += 1
rec("2Q26 SanDisk 營收（實績重現）", "=rv_sd_5", lambda s: f"=Mkt!{QCOL['2026Q2']}{MB[s]['rv_sd']}", TF0818, "恆等式", "#,##0.00")
rec("2026Q3 SanDisk 營收 vs 公司指引（TrendForce 口徑）", "=sd_g/sd_k", lambda s: f"=Mkt!{QCOL['2026Q3']}{MB[s]['rv_sd']}", "SanDisk 指引 $10.3–10.8B", "恆等式（份額由指引推得）", "#,##0.00")
rec("2026Q3 市場營收 vs 2Q26 ×（1＋季價 15%）", "=Mkt!" + QCOL["2026Q2"] + str(MB["A"]["mq"]) + "*1.15", lambda s: f"=Mkt!{QCOL['2026Q3']}{MB[s]['mq']}", "TrendForce 3Q26 合約價 +10~15%（markets/memory）", "差額＝位元季增")
rec("CY2026 位元出貨年增 vs 美光「低 20% 段」（百分點）", 0.21, lambda s: f"=Mkt!{YC[2026]}{MB[s]['bg']}", "美光 FQ4 FY26（Guru3D 2026-10-01）", "供給受限時＝gs_26", "0.0%", True)
rec("CY2027 位元出貨年增 vs 美光「約中 20% 段」（百分點）", 0.25, lambda s: f"=Mkt!{YC[2027]}{MB[s]['bg']}", "同上", "", "0.0%", True)
rec("CY2028 位元出貨年增 vs 美光「約中 20% 段」（百分點）", 0.25, lambda s: f"=Mkt!{YC[2028]}{MB[s]['bg']}", "同上", "B、C 需求低於供給，出貨＝需求", "0.0%", True)
rec("CY2026 需求 ÷ 供給 − 1 vs TrendForce 供給缺口 4–6%（百分點）", 0.05, lambda s: f"=Mkt!{YC[2026]}{MB[s]['gap']}", "TrendForce 2026-07-21（Guru3D 轉載）：2026 供給低於需求 4–6%；2H27 回到平衡", "", "0.0%", True)
rec("CY2026 企業級 SSD 占需求 vs Counterpoint 2Q26 48%（百分點）", 0.48, lambda s: f"=Mkt!{YC[2026]}{MB[s]['ess']}", "Counterpoint 2Q26（季度值；全年平均應較低）", "TrendForce：2026 伺服器約 40%", "0.0%", True)
rec("CY2025 市場營收（實績）vs 前五大合計 ÷ 0.9", "=(" + "+".join(f"rv_{k}_{i}" for k in TOP5 for i in range(4)) + ")/0.9", lambda s: f"=Mkt!{YC[2025]}{MB[s]['rev']}", "TrendForce 季度合計", "")
rec("CY2027 Kioxia＋SanDisk 營收 vs 共識（會計年度依月份拼成日曆年：Kioxia ¼ FY2026＋¾ FY2027；SanDisk ½ FY2027＋½ FY2028）", "=10071/160/kx_k*0.25+13477/160/kx_k*0.75+49.1/sd_k*0.5+58.4/sd_k*0.5",
    lambda s: f"=Ann!{YC[2027]}{AR[(s, 'kxr')]}+Ann!{YC[2027]}{AR[(s, 'sdr')]}", "MarketScreener 2026-10-23 擷取：Kioxia FY2027（2027-03）營收 ¥10.07 兆、FY2028 ¥13.48 兆；SanDisk FY2027 $49.1B、FY2028 $58.4B", "會計年度依月份拼成日曆年（Derived）；共識只作對照")
RC_LAST = rcr - 1
RC.column_dimensions["A"].width = 70; RC.column_dimensions["G"].width = 60; RC.column_dimensions["H"].width = 40
for c in "BCDEF": RC.column_dimensions[c].width = 11

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A）：對 CY2027、CY2028 市場營收與 Kioxia＋SanDisk 營收的影響。每個案例為一個完整推導區塊（本頁下方，與 Mkt 同一引擎）"
hdr(SE, 3, ["案例", "設定", "CY2027 市場 $B", "相對基準", "CY2028 市場 $B", "相對基準", "CY2028 位元出貨 EB", "CY2028 合資群營收 $B", "相對基準", "CY2030 市場 $B", "相對基準"])
CASES = [("每顆容量 −20%", "tb 低欄", dict(tb="_lo")), ("每顆容量 +20%", "tb 高欄", dict(tb="_hi")),
         ("位元供給成長 低", "gs 低欄", dict(gs="_lo")), ("位元供給成長 高", "gs 高欄", dict(gs="_hi")),
         ("2025 位元基期 850 EB", "bits25 低", dict(bits="bits25_lo")), ("2025 位元基期 1,150 EB", "bits25 高", dict(bits="bits25_hi")),
         ("企業級 SSD 2025 占比 26%", "sh_essd25 低", dict(essd="sh_essd25_lo")), ("非 AI 企業級 SSD 年增 5%", "g_nai 低", dict(nai="g_nai_lo")),
         ("2029 顆數年增 0%", "g_u29 低", dict(u29="g_u29_lo")), ("CY2030 售價年變動 −20%", "p10 低", dict(p10="p10_lo")),
         ("2027 起季價水準 ×0.8", "價格係數 0.8", dict(px="0.8")), ("合資群份額年變動取 C（−0.6pt）", "dr_jv_C", dict(jv="dr_jv_C"))]
def put(i, nm, st, R, pre=""):
    SE.cell(i, 1, nm); SE.cell(i, 2, st)
    SE[f"C{i}"] = f"={pre}{YC[2027]}{R['rev']}"; SE[f"E{i}"] = f"={pre}{YC[2028]}{R['rev']}"; SE[f"G{i}"] = f"={pre}{YC[2028]}{R['ship']}"
    SE[f"H{i}"] = f"=SUM({pre}{QCOL['2028Q1']}{R['rv_kx']}:{QCOL['2028Q4']}{R['rv_kx']})+SUM({pre}{QCOL['2028Q1']}{R['rv_sd']}:{QCOL['2028Q4']}{R['rv_sd']})"
    SE[f"D{i}"] = f"=C{i}/C$4-1"; SE[f"F{i}"] = f"=E{i}/E$4-1"; SE[f"I{i}"] = f"=H{i}/H$4-1"
    SE[f"J{i}"] = f"={pre}{YC[2030]}{R['rev']}"; SE[f"K{i}"] = f"=J{i}/J$4-1"
    for c in "CEGHJ": SE[f"{c}{i}"].number_format = "#,##0.0"
    for c in "DFIK": SE[f"{c}{i}"].number_format = "0.0%"
put(4, "基準（情境 A）", "Mkt 區塊 A", MB["A"], "Mkt!")
base = 40
for n, (nm, st, ov) in enumerate(CASES):
    Rn = engine(SE, base + n * 70, f"案例區塊：{nm}（情境 A）", "A", **ov)
    put(5 + n, nm, st, Rn)
SN = 5 + len(CASES)
put(SN, "情境 B", "Mkt 區塊 B", MB["B"], "Mkt!"); put(SN + 1, "情境 C", "Mkt 區塊 C", MB["C"], "Mkt!")
SE.cell(SN + 3, 1, "讀法：營收以 2Q26 實績為錨，2026Q3 起的變動＝價格指數比 × 位元比；Kioxia、SanDisk 的 2026Q3 份額由公司指引推得，因此改變市場路徑時兩家 2026Q3 營收不變（份額反向調整）。")
SE.column_dimensions["A"].width = 40; SE.column_dimensions["B"].width = 26
for c in "CDEFGHI": SE.column_dimensions[c].width = 13

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源"; SO["A1"].font = BOLD
srcs = [("markets/memory/market.py：NAND 季價路徑（NAND_COMMON、NAND_S）、加速器顆數 UNITS、情境顆數係數 SCEN_UNITS", "本 repo（單一來源）"),
        ("TrendForce：2Q26 NAND 前五大營收（2026-08-18）", "https://www.trendforce.com/presscenter/news/20260818-13186.html"),
        ("TrendForce：2Q25 NAND 營收（2025-08-28）", "https://www.trendforce.com/presscenter/news/20250828-12688.html"),
        ("TrendForce：3Q25 NAND 營收（2025-12-03）", "https://www.trendforce.com/presscenter/news/20251203-12813.html"),
        ("TrendForce：4Q25 NAND 營收（I-Connect007 轉載）", "https://iconnect007.com/article/149024/ai-server-storage-demand-surges-top-five-nand-flash-suppliers-post-238-qoq-revenue-growth-in-4q25/149021/pcb"),
        ("TrendForce：2024 NAND 營收 $67.4B（TelecomLead 轉載）", "https://telecomlead.com/semiconductor/memory-industry-to-see-significant-revenue-growth-trendforce-117301"),
        ("TrendForce：2026 供給缺口 4–6%、2H27 平衡、伺服器約 40%（Guru3D 2026-07-21 轉載）", "https://www.guru3d.com/story/nand-flash-supply-expected-to-return-to-market-balance-in-2027/"),
        ("美光：產業 NAND 位元出貨 CY2026–28（Guru3D 2026-10-01 轉載 FQ4 FY26）", "https://www.guru3d.com/story/micron-expects-dram-and-nand-supply-to-tighten-further-through-2028/"),
        ("Counterpoint：2Q26 企業級 SSD 占位元 48%、位元市占（capitalandcompute、hwbusters 轉載）", "https://capitalandcompute.net/blog/why-are-ssds-so-expensive/"),
        ("SanDisk FQ4 FY26 新聞稿（8-K，2026-08-05）", "https://www.sec.gov/Archives/edgar/data/0002023554/000162828026053346/sndkq4-26ex991xpressrelease.htm"),
        ("Kioxia FY2026 第一季財報（SGX 轉載，2026-07-31）", "https://links.sgx.com/1.0.0/corporate-announcements/DOATO9EIE4EEQ2Y8/898078_Earnings-Release-FY2026-1Q-EN.pdf"),
        ("borecraft：Kioxia 332 層、三星／SK 投片（2026-05-26，轉載 DigiTimes 等；AI 輔助撰寫，弱）", "https://borecraft.com/?p=15672"),
        ("TrendForce：Kioxia FY26–28 資本支出（2026-06-03）", "https://www.trendforce.com/news/2026/06/03/news-kioxia-targets-%C2%A5470b-yearly-capex-in-fy26-28-up-66-from-fy25-reportedly-weighing-third-kitakami-fab-and-ma/"),
        ("MarketScreener：Kioxia、SanDisk 共識（2026-10-23 擷取）", "https://uk.marketscreener.com/quote/stock/KIOXIA-HOLDINGS-CORPORATI-180030735/finances/"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 120

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("NAND 市場層 v0.1（AI 半導體 Project；樣板 markets/memory）　2026-10-25", True),
 ("命題：NAND 營收由 AI 推論帶動的企業級 SSD 位元需求與合約價循環決定；2026–27 供給受限、價格在高檔，2028 起各情境依加速器顆數係數與價格路徑分歧。", False), ("", False),
 ("驅動因子 → 推導量", True),
 ("Mkt A 區｜加速器顆數（markets/memory UNITS × 同字母 SCEN_UNITS）× 每顆企業級 SSD 容量 → AI 企業級 SSD 需求；非 AI 企業級 SSD、手機、PC、消費性以 2025 基期 × 年增 → 位元需求合計", False),
 ("Mkt A 區｜全產業位元供給（2025 基期 ×（1＋美光指引年增），三情境相同）；出貨＝MIN(需求, 供給)", False),
 ("Mkt B 區｜季價指數（markets/memory NAND 路徑，同字母；2030 收斂法）× 季度位元 → 以 2Q26 實績為錨的季度市場營收；份額（2026Q3 由 Kioxia、SanDisk 公司指引推得）→ 各廠營收與位元", False),
 ("Ann｜年度摘要：需求、供給、出貨、平均售價、營收、各廠份額；Kioxia＋SanDisk 合資群份額＝兩家相加（下游公司模型讀同一組列）", False), ("", False),
 ("情境：A＝順利、B＝平平、C＝不如預期；價格讀 markets/memory NAND_S 同字母，顆數讀 SCEN_UNITS 同字母（2028–29；CY2030 沿用 2029），不自訂行業路徑。CY2030 價格為收斂法（十年平均年變動 p10，三情境相同）。", False),
 ("單位：$B、EB、$/GB、百萬顆。2026Q2 以前為 TrendForce 實績；季度位元為年度依年增平滑（近似）。", False),
 ("已知限制：全產業 EB 基期與各應用占比為 Assumed（只影響 AI 與非 AI 需求的相對大小與 EB 絕對值）；各廠晶圓產能未找到（Sup 頁）；各廠平均售價假設等於市場平均。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 180
order = ["README", "In", "Ann", "Path", "Mkt", "Sup", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
print("saved", sys.argv[1])
