# 半導體設備市場模型 v0.1.2（AI 半導體 Project 步驟 5：設備與材料——設備部分）
# 用法：python3 markets/equipment/build_equipment.py raw.xlsx ；再以 LibreOffice headless 重算
# 需求＝為滿足記憶體、代工擴產所需的資本支出：資本支出讀自 skhynix／micron／samsung／tsmc 的 dist 成品（不修改那些模型）
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

BLUE = Font(color="0000FF"); GREEN = Font(color="007A33"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(wrap_text=True, vertical="top")
YRS = (25, 26, 27, 28, 29); YC = {25: "C", 26: "D", 27: "E", 28: "F", 29: "G"}

# ---------------- 讀四家公司 dist ----------------
def dist_file(folder):
    fs = glob.glob(os.path.join(ROOT, folder, "dist", "*.xlsx"))
    if len(fs) != 1: raise SystemExit(f"{folder}/dist 應剛好一個 xlsx，實際 {len(fs)}")
    return fs[0]

def read_capex(folder, sheet, label, heads):
    """讀 sheet 內 A 欄＝label 那列；heads＝2026–2029 的 A 情境欄名；B、C 情境欄必須與 A 相同（資本支出不隨情境變）。"""
    f = dist_file(folder); wb = openpyxl.load_workbook(f, data_only=True); ws = wb[sheet]
    hdr = {ws.cell(3, j).value: j for j in range(1, ws.max_column + 1) if ws.cell(3, j).value}
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == label)
    out = {}
    for y, h in zip((26, 27, 28, 29), heads):
        j = hdr[h]; v = ws.cell(row, j).value
        if y > 26:
            for s in "BC":
                jj = hdr[h.replace(" A", f" {s}")]
                assert abs(ws.cell(row, jj).value - v) < 1e-9, f"{folder} 資本支出隨情境不同"
        out[y] = (v, f"{os.path.basename(f)}／{sheet}!{openpyxl.utils.get_column_letter(j)}{row}")
    return out, wb

def in_row(wb, key):
    ws = wb["In"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == key: return [ws.cell(r, j).value for j in range(1, 10)]
    raise SystemExit(f"找不到 {key}")

SK, wb_sk = read_capex("skhynix", "Y", "資本支出", ["2026（基準）", "2027 A", "2028 A", "2029 A"])
MU, wb_mu = read_capex("micron", "Y", "資本支出", ["FY2026（基準）", "FY2027 A", "FY2028 A", "FY2029 A"])
SS, wb_ss = read_capex("samsung", "Y", "資本支出", ["2026（基準）", "2027 A", "2028 A", "2029 A"])
TS, wb_ts = read_capex("tsmc", "PL", "資本支出", ["2026", "2027 A", "2028 A", "2029 A"])
FX = in_row(wb_sk, "fx")            # 低／基準／高（Derived，SK 海力士模型）
CAPSRC = {  # 公司模型 In 頁該年資本支出的標記與來源（第二來源）
    "sk": [in_row(wb_sk, k) for k in ("cap_0", "cap_1", "cap_2", "cap_3")],
    "mu": [in_row(wb_mu, k) for k in ("cap_0", "cap_1", "cap_2", "cap_3")],
    "ss": [in_row(wb_ss, k) for k in ("cap_0", "cap_1", "cap_2", "cap_3")],
    "ts": [in_row(wb_ts, k) for k in ("capex26", "capex27", "capex28", "capex29")]}

wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字＝本模型輸入；綠字＝讀自其他模型 dist 成品，不在此修改；基準欄 E）。金額：十億美元（公司資本支出另標幣別）"
IN["A2"] = "每個數字附兩個來源（來源1、來源2）；找不到第二來源者寫「未找到」。SemiAnalysis 未使用。標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／備註"])
rows = []
sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)

SEMI_URL = "SEMI 2026-07-14 年中預測新聞稿 https://www.semi.org/en/semi-press-release/global-semiconductor-equipment-sales-forecast-to-reach-a-record-229-billion-dollars-in-2028-semi-reports"
SEMI_2 = "轉載：PR Newswire／Semiconductor Digest（同一原始發布，非獨立來源）；SEMI 為設備商協會，傾向樂觀（Interested-party）"
sec("A. 公司資本支出（2026–2029 讀自 dist 成品，三情境相同；2025 為本模型自行輸入——其他模型不含 2025——用於求「其他廠」2025 基準）")
inp("fx", "美元兌韓元（2025 亦沿用，Assumed）", "KRW/USD", FX[3], FX[4], FX[5], "Derived", "讀自 skhynix 模型 In!fx（" + str(FX[7]) + "）", "三星模型 In!fx 同值（「同 SK 海力士」）；獨立匯率來源：本模型未另查")
for y, (v, ref) in SK.items():
    r = CAPSRC["sk"][y - 26]
    inp(f"cap_sk_{y}", f"SK 海力士資本支出 {2000+y}", "兆韓元", None, v, None, "Interested-party" if y < 29 else "Assumed", ref, f"模型 In 標記：{r[6]}｜{r[7]}｜第二來源：未找到（模型只引共識）" if y < 29 else f"模型 In 標記：{r[6]}｜{r[7]}｜第二來源：未找到")
for y, (v, ref) in MU.items():
    r = CAPSRC["mu"][y - 26]
    inp(f"cap_mu_{y}", f"美光資本支出 FY{2000+y}（淨；以 FY≈日曆年）", "$B", None, v, None, r[6], ref, f"模型 In 標記：{r[6]}｜{r[7]}｜美光 FY 於 8–9 月結束，以日曆年近似，未調整（誤差方向：FY 比日曆年早約 4 個月）")
for y, (v, ref) in SS.items():
    r = CAPSRC["ss"][y - 26]
    inp(f"cap_ss_{y}", f"三星記憶體資本支出 {2000+y}（只取記憶體）", "兆韓元", None, v, None, r[6], ref, f"口徑（三星模型 In）：{r[6]}｜{r[7]}；全公司×記憶體約 75% 為 Assumed；第二來源：未找到")
for y, (v, ref) in TS.items():
    r = CAPSRC["ts"][y - 26]
    inp(f"cap_ts_{y}", f"台積電資本支出 {2000+y}", "$B", None, v, None, r[6], ref, f"模型 In 標記：{r[6]}｜{r[7]}")

inp("cap_sk_25", "SK 海力士資本支出 2025", "兆韓元", 28, 30.2, 30.2, "Interested-party", "Qz 2026-07-29（公司在 2Q26 法說稱 2026 預算相較 2025 的 30.2 兆韓元）https://qz.com/sk-hynix-q2-2026-earnings-record-profit-misses-estimates-072926", "MLQ https://mlq.ai/news/sk-hynix-posts-record-605-trillion-won-operating-profit-raises-capacity-spending/ 同為 2Q26 法說摘要；另有產業估計約 28 兆（X：Jukan，低 28）。皆為搜尋摘要，未開啟原文")
inp("cap_mu_25", "美光資本支出 FY2025（淨）", "$B", None, 13.80, None, "Verified", "美光 FY25 8-K（淨 13.80）https://www.sec.gov/Archives/edgar/data/723125/000072312525000024/a2025q4ex991-pressrelease.htm", "美光 FY25 10-K（毛 15.86、政府補助 2.01）https://www.sec.gov/Archives/edgar/data/723125/000072312525000028/mu-20250828.htm")
inp("cap_ssds_25", "三星 DS（半導體）資本支出 2025", "兆韓元", None, 47.5, None, "Verified", "三星 IR：FY2025 CAPEX Results（DS 47.5、顯示 2.8、合計約 52.7）https://www.samsung.com/global/ir/reports-disclosures/public-disclosure-view.84547/", "Korea Herald 合計 52.7 兆（https://koreaherald.com/article/10691933）；DS 47.5 另見 tech-insider（低權威）")
inp("ss_mem", "三星 DS 資本支出中記憶體占比（2025）", "%", 0.65, 0.75, 0.90, "Assumed", "沿用三星模型 In（全公司 × 記憶體約 75%，Assumed）；三星不揭露記憶體分項。區間上修到 90%：2025 年三星晶圓代工投資據報減半至約 5 兆韓元（2024 約 10 兆；Sedaily 轉引，https://www.trendforce.com/news/2025/01/22/news-samsung-faces-struggles-ahead-as-foundry-investment-reportedly-slashed-by-half-for-2025/；https://www.digitimes.com/news/a20250122PD229/samsung-capex-business-plant-investment.html），記憶體占比較可能高於 75%", "未找到記憶體分項的第二來源。算術參考（非來源）：DS 47.5 兆扣除晶圓代工約 5 兆、LSI 未知，記憶體約 85–89%；基準按指示維持 75%。此假設直接影響 2025 三大廠資本支出")
inp("cap_ts_25", "台積電資本支出 2025", "$B", 40, 40.9, 42, "Interested-party", "Semiconductor Intelligence：台積電 2025 資本支出 40.9B https://www.semiconductorintelligence.com/capex-up-for-foundry-memory/", "台積電 3Q25 法說：全年指引 40–42B（https://www.fool.com/earnings/call-transcripts/2025/10/17/tsmc-tsm-q3-2025-earnings-call-transcript）")
sec("B. 設備占資本支出比例：按公司分開（公司有揭露建築與設備比例者用公司揭露；未找到者用業界 58–70%，Analogy）")
SH_SRC1 = "SEMI 2025 WFE 116.9B（年中預測新聞稿）÷ Semiconductor Intelligence 2025 全球半導體資本支出 166B（https://www.semiconductorintelligence.com/capex-up-for-foundry-memory/）＝70.4% → 高 70%"
SH_SRC2 = "Rebound Capital（Substack）資本支出→晶圓前段設備轉換率約 58%（https://reboundcapital.substack.com/p/industry-deep-dive-semiconductor）→ 低 58%；基準取兩者中點 64%。業界合計值，未找到記憶體／代工分項"
# 美光：由公司揭露推導（建築占增量比例 c；FY25 基準設備占比＝業界中點）
MU25, MU26, MU27 = 13.80, MU[26][0], MU[27][0]; B25 = 0.64
def mu_share(cap_y, cap_prev, cons_prev, c): return 1 - (cons_prev + c * (cap_y - cap_prev)) / cap_y
cons25 = (1 - B25) * MU25
SHM = {}
for c_, tag in ((0.5, "hi"), (0.6, "base"), (0.7, "lo")):
    cons26 = cons25 + c_ * (MU26 - MU25); cons27 = cons26 + c_ * (MU27 - MU26)
    SHM[tag] = (round(1 - cons26 / MU26, 3), round(1 - cons27 / MU27, 3))
MU_Q3 = "FQ3 FY26 準備稿（https://s25.q4cdn.com/621799436/files/doc_financials/2026/q3/Q3-FY26-Prepared-Remarks.pdf）：「more than half the increase year over year in fiscal 2027 from construction capex as we pull in cleanroom capacity required to address long-term demand」（引自搜尋摘要，未能開啟全文）"
MU_Q4 = "FQ4 FY26 法說（https://s25.q4cdn.com/621799436/files/doc_financials/2026/q4/Q4-FY26-Prepared-Remarks.pdf；逐字稿 https://www.fool.com/earnings/call-transcripts/2026/10/01/micron-mu-q4-2026-earnings-call-transcript/）：「A majority of the increase is for construction capex, most of which is to help accelerate clean room space availability in late calendar 2028 and beyond」「a meaningfully higher growth rate in construction CapEx as compared to equipment CapEx in fiscal 2027」（引自搜尋摘要）"
MU_NOTE = ("注意：公司原文是「資本支出『增量』多數為建築」，不是「FY27 資本支出多數為建築」，故文稿本身不給出設備占比，只給出下限性的建築比例。"
           "推導：FY25 淨資本支出 13.80（8-K）× (1−64%) ＝建築基數；之後每年增量中建築占 c（c＝50%「多於一半」為上限情境，60%、70% 為 Assumed 區間）；"
           f"FY26＝{MU26:.2f}、FY27＝{MU27:.0f}（美光模型）。FY28–FY29 無揭露，沿用 FY27。")
inp("sh_mu_25", "美光：FY25 設備占資本支出比例（建築基數的起點）", "%", 0.58, 0.64, 0.70, "Analogy", SH_SRC1, SH_SRC2 + "。美光未揭露 FY25 建築／設備比例")
inp("sh_mu_26", "美光：FY26 設備占資本支出比例", "%", SHM["lo"][0], SHM["base"][0], SHM["hi"][0], "Derived", MU_Q3 + "｜" + MU_Q4 + "｜FY26 增量多數為無塵室建築（摘要轉述，非逐字）", MU_NOTE + " 引文為公司一手（Verified），比例為 Derived（FY25 基準 64% 為 Analogy、c 為 Assumed）")
inp("sh_mu_27", "美光：FY27–FY29 設備占資本支出比例", "%", SHM["lo"][1], SHM["base"][1], SHM["hi"][1], "Derived", MU_Q3 + "；" + MU_Q4, MU_NOTE + " 高＝c 取 50%（「多於一半」）時的推導上限，約 53%。引文為公司一手（Verified），比例為 Derived")
inp("sh_sk", "SK 海力士：設備占資本支出比例（2025 實績校準值）", "%", 0.0, 0.0, 0.0, "Derived", "Mkt 頁公式：令 2025 年「SK＋三星＋美光設備＋其他廠獨立估計」＝SEMI 2025 DRAM＋NAND 實績（38.5），解出 SK、三星共用的一個占比；2026 年以後沿用。低＝長鑫取 12.0 時的校準值；高＝58%（業界下限，Rebound Capital）", "校準只用 2025 實績；2026 年以後完全沒有參照 SEMI 或券商的預測。未找到 SK 海力士公司揭露（定性：董事會 2026-02 核准 21.6 兆韓元完成龍仁第一晶圓廠結構工程與二至六期無塵室）。業界值 64% 不再當基準，Sens 保留一列對照")
inp("sh_ss", "三星（記憶體）：設備占資本支出比例（同 SK 海力士，校準值）", "%", 0.0, 0.0, 0.0, "Derived", "同 sh_sk（兩者共用同一個校準值，以 2025 實績校準）", "未找到三星公司揭露（P5 無塵室先行，無建築／設備比例）。三星記憶體 2025 資本支出為 DS 47.5 兆 × ss_mem（Assumed）；ss_mem 改變時校準值隨之改變（Sens）")
inp("sh_fd", "台積電（代工與邏輯）：設備占資本支出比例", "%", 0.50, 0.64, 0.75, "Analogy", SH_SRC1, SH_SRC2 + "。v0.1.2 查證（TSMC 頁）：台積電董事會每季資本預算有分「機器設備」與「不動產與資本化租賃資產」（二手摘要，5 次，彙總設備占 60%，單次 47–71%）；元大投顧稱先進製程外商前段設備占資本支出 7～8 成（Interested-party）。兩者都未顯示台積電像 SK 海力士、三星（45.5%）那樣偏低，故基準維持 64%；低 50%＝董事會 2025-08（47%）、2026-02（52%）兩次的水準，高 75%＝元大 7～8 成的中點。年報不動產廠房及設備附註的分類增添：未找到（本環境無法開啟年報 PDF）。台積電資本支出含先進封裝、光罩等（公司稱 10–20%），base 未另扣，見 pk_ts 與 Sens")
sec("B1b. 台積電設備占比證據（v0.1.2；明細見 TSMC 頁）")
inp("pk_ts", "台積電資本支出中屬先進封裝、測試、光罩及其他的比例", "%", 0.10, 0.15, 0.20, "Interested-party", "台積電（3Q25 法說逐字稿，Motley Fool）：2025 資本支出約 10–20% 用於先進封裝、測試、光罩及其他 https://www.fool.com/earnings/call-transcripts/2025/10/17/tsmc-tsm-q3-2025-earnings-call-transcript", "2026 年配置同為 10–20%（搜尋摘要：https://xmy.tw/articles/global-semiconductor-capex-2026/tsmc-capex-2026，未開啟原文）。公司揭露的是「用途」分類，不是設備／建築分類；其設備中有多少被 SEMI 分入後段（測試、組裝封裝）或其他前段（光罩）：未找到。base 不扣，只進 Sens 與 TSMC 頁")
inp("cap_smic_25", "中芯 2025 資本支出（TSMC 頁可識別性檢查用）", "$B", 8.1, 8.1, 8.1, "Interested-party", "TrendForce（2026-03-30）：中芯 2025 約 $8.1B、2026 約持平 https://www.trendforce.com/news/2026/03/30/news-smic-2026-action-plan-points-to-above-industry-growth-capex-held-in-line-with-2025/", "第二來源：未找到。沿用本模型 g_oth_fd 說明所引數字")
inp("cap_intel_25", "Intel 2025 資本支出（毛額；扣抵後淨額約 11.2）", "$B", 11.2, 17.7, 17.7, "Interested-party", "Intel 4Q25 法說摘要（Yahoo Finance／MarketBeat）：毛額 $17.7B、扣抵約 $6.5B → 淨約 $11.2B（Derived）https://finance.yahoo.com/news/intel-q4-earnings-call-highlights-234233695.html", "第二來源：Intel FY2025 10-K（https://www.sec.gov/Archives/edgar/data/50863/000005086326000011/intc-20251227.htm，搜尋結果列出，未開啟核對）；第三方資料站現金流量口徑 14.65（口徑不同）。SEMI 計的是設備出貨，取毛額為基準")
sec("B2. 其他廠 2026 年以後成長率（Assumed；每個區塊合併成一個年成長率，套用 2026–2029。依據公司資本支出指引；不參照 SEMI 2026 以後預測）")
inp("g_oth_dram", "其他 DRAM 廠（長鑫、南亞、華邦等）設備成長率", "%/年", 0.10, 0.18, 0.60, "Assumed", "長鑫：未揭露資本支出；研究估計月產能 360k／420k／500k 片（2026／27／28）→ +17%／+19%（https://faxiangongchang.com/en/reports/china-memory-chip-2026，弱）；南亞 2026 資本支出上限 NT$69.7B，2027 最高 NT$200B（約 4 倍，https://www.digitimes.com/news/a20260713PD211/nanya-technology-capex-capacity-2027.html）；華邦 2026 NT$42.1B（紀錄高，https://www.gate.com/news/detail/winbond-approves-507b-capex-expands-kaohsiung-fab-production-to-24000-23369791）", "基準取長鑫產能成長（兩年平均約 18%，Analogy：產能成長代替資本支出）；低 10% 假設資本支出增速低於產能；高 60% 取南亞 2027 與鎧俠 FY26 的高增速。長鑫的資本支出指引：未找到；2025 年基準（華邦、南亞）：未找到")
inp("g_oth_nand", "其他 NAND 廠（鎧俠、SanDisk 合資、長江存儲等）設備成長率", "%/年", 0.0, 0.18, 0.60, "Assumed", "鎧俠：FY25 ¥283.7B → FY26 ¥450B（+59%），FY26–28 年均約 ¥470B（https://www.trendforce.com/news/2026/06/03/news-kioxia-targets-%C2%A5470b-yearly-capex-in-fy26-28-up-66-from-fy25-reportedly-weighing-third-kitakami-fab-and-ma/）；SanDisk FY27 毛資本支出約營收 6%（含與鎧俠合資的分攤，與鎧俠重疊，不另加，https://www.fool.com/investing/2026/08/29/sandisk-and-kioxia-plan-to-invest-more-than-usd31-billion-in-japanese-memory-plants-about-60-of-what-they-have-spent-there-in-25-years/）；長江存儲：三廠 2026 底投產，金額未找到", "基準＝鎧俠 FY25→FY28 年複合 ((470/283.7)^(1/3)−1)≈18%（Derived，單一公司）；低 0% 為 FY27–28 持平；高 60% 為 FY26 單年增速。長江存儲資本支出指引：未找到")
inp("g_oth_fd", "其他代工與邏輯（中芯、Intel、三星晶圓代工、華虹、聯電、GlobalFoundries 等）設備成長率", "%/年", 0.0, 0.08, 0.20, "Assumed", "中芯：2025 約 $8.1B、2026 約持平（https://www.trendforce.com/news/2026/03/30/news-smic-2026-action-plan-points-to-above-industry-growth-capex-held-in-line-with-2025/）；Intel：2026 提高到 >$20B（原約 $17–18B；2027「明顯更高」，https://www.gate.com/news/detail/intel-raises-2026-capex-guidance-to-over-20-billion-q2-earnings-beat-22949346）", "基準＝中芯（0%）與 Intel（約 +11%，18→20）以 2025 金額加權約 +8%（Derived）；低 0%＝中芯持平；高 20% 假設 Intel 2027 大增。三星晶圓代工、華虹、聯電、GlobalFoundries 資本支出指引：未找到（三星僅有半導體合計）")
sec("G. 2025 年其他記憶體廠資本支出：獨立估計（用作 2025 其他廠基準，並用於校準 SK 海力士、三星設備占比；只用 2025 實績）")
inp("sh_oth", "其他記憶體廠設備占資本支出比例（業界值）", "%", 0.58, 0.64, 0.70, "Analogy", SH_SRC1, SH_SRC2)
inp("ind_cxmt", "長鑫 2025 資本支出（高值年度未核對）", "$B", 6.9, 6.9, 12.0, "Interested-party", "TechInsights（Reuters 轉引，2025 初）：2023、2024 約 $6–7B，預期 2025 +5%（≈6.9）https://finance.yahoo.com/news/exclusive-chipmaker-cxmt-plans-shanghai-073043365.html", "高 12.0：BigGo 摘要「CNY 80.5B（約 $12.0B）、+62%」，年度未核對 https://finance.biggo.com/news/5a548f95-479d-4e07-8dd1-af5beb8eea1f。TrendForce、券商的長鑫 2025 資本支出：未找到。基準取低端，無依據偏好")
inp("ind_nand_ind", "全球 NAND 資本支出 2025", "$B", None, 21.1, None, "Interested-party", "TrendForce 2025-11-13：NAND 資本支出 2025 $21.1B、2026 $22.2B https://www.trendforce.com/presscenter/news/20251113-12780.html", "StorageNewsletter 轉載 https://www.storagenewsletter.com/2025/12/31/memory-industry-to-maintain-cautious-capex-in-2026/（非獨立）")
inp("ind_ymtc_sh", "長江存儲占全球 NAND 資本支出比例（2025）", "%", None, 0.20, None, "Interested-party", "搜尋摘要稱 YMTC 2025 資本支出約占全球 NAND 投資 20%（出處為 Digitimes 系列報導，確切文章未核對）", "第二來源：未找到")
inp("ind_kioxia", "鎧俠 FY2025 資本支出（FY 至 2026-03）", "¥B", None, 283.7, None, "Interested-party", "TrendForce 2026-06-03（FY26 ¥450B、FY26–28 年均 ¥470B，較 FY25 +66%；FY25 ¥283.7B 為二手整理）https://www.trendforce.com/news/2026/06/03/news-kioxia-targets-%C2%A5470b-yearly-capex-in-fy26-28-up-66-from-fy25-reportedly-weighing-third-kitakami-fab-and-ma/", "鎧俠 Investor Day 2026 簡報 https://www.kioxia-holdings.com/content/dam/kioxia-hd/en-jp/ir/library/event/asset/Kioxia-Investor-Day-2026-en.pdf（未開啟，未核對 ¥283.7B）")
inp("fx_jpy", "日圓兌美元（2025）", "JPY/USD", 140, 150, 155, "Assumed", "無：未查匯率來源（2025 年大致在 140–155）", "—")
inp("ind_nanya", "南亞科 2025 資本支出", "NT$B", None, 13.4, None, "Verified", "南亞科 IR：2025 預算上限 19.6、實支 13.4、6.2 遞延到 2026 https://www.nanya.com/en/IR/16/?IRId=11068", "第三方 Q1 簡報摘要同值（見 IR 頁）")
inp("fx_twd", "新台幣兌美元（2025）", "TWD/USD", 29.5, 31, 32.5, "Assumed", "無：未查匯率來源", "—")
inp("ind_winbond", "華邦 2025 資本支出", "NT$B", None, None, None, "—（未找到）", "未找到（只查到 2026 年 NT$42.1B，https://www.gate.com/news/detail/winbond-approves-507b-capex-expands-kaohsiung-fab-production-to-24000-23369791）", "未找到；估計記為 0（總額偏低）")
inp("ind_sandisk", "SanDisk 2025 資本支出（JV 分攤）", "$B", None, None, None, "—（未找到）", "FY25 數字未找到；FY26 毛資本支出 $0.562B，其中 JV 分攤 $0.519B（Q4 FY26 簡報 https://investor.sandisk.com/static-files/c75d1bee-c5c9-4e5a-8605-302c1aeac59b）", "JV 分攤與鎧俠重疊（Flash Ventures 由鎧俠 50.1％、SanDisk 49.9％持有），不另計；非 JV 部分約 $0.04B（FY26），忽略")
sec("C. SEMI 2026-07 年中預測（Interested-party）。2025 為實績（由年增率回推），用於求其他廠基準；2026 以後的 SEMI 分項只用於對帳，不決定本模型總額（測試與封裝除外，v0.1 以 SEMI 為輸入）")
for nm, k, lvl, g26, g27, g28 in (("DRAM 設備", "dram", 38.8, 0.390, 0.274, 0.150), ("NAND 設備", "nand", 13.9, 0.307, 0.311, 0.145),
                                ("代工與邏輯設備", "fd", 78.0, 0.189, 0.181, 0.136), ("晶圓前段設備合計（WFE）", "wfe", 143.9, 0.231, 0.218, 0.141)):
    inp(f"s_{k}_26", f"SEMI {nm} 2026", "$B", None, lvl, None, "Interested-party", SEMI_URL, SEMI_2)
    inp(f"g_{k}_26", f"SEMI {nm} 2026 年增", "%", None, g26, None, "Interested-party", SEMI_URL, SEMI_2)
    inp(f"g_{k}_27", f"SEMI {nm} 2027 年增", "%", None, g27, None, "Interested-party", SEMI_URL + "（2027 前段 +21.8% 與其餘分項年增皆來自 SEMI 新聞稿摘要）", SEMI_2)
    inp(f"g_{k}_28", f"SEMI {nm} 2028 年增", "%", None, g28, None, "Interested-party", SEMI_URL, SEMI_2 + "；DRAM 2028 推得 56.85 vs 報導 56.9、NAND 20.87 vs 20.8（捨入）")
inp("s_test_26", "SEMI 測試設備 2026", "$B", None, 15.3, None, "Interested-party", SEMI_URL, SEMI_2)
inp("g_test_26", "SEMI 測試設備 2026 年增", "%", None, 0.310, None, "Interested-party", SEMI_URL, SEMI_2)
inp("s_test_28", "SEMI 測試設備 2028", "$B", None, 20.8, None, "Interested-party", SEMI_URL, SEMI_2 + "。2027 年值未找到 → Mkt 以 2026、2028 幾何內插（Derived）")
inp("s_pkg_26", "SEMI 組裝與封裝設備 2026", "$B", None, 6.7, None, "Interested-party", SEMI_URL, SEMI_2)
inp("g_pkg_26", "SEMI 組裝與封裝設備 2026 年增", "%", None, 0.096, None, "Interested-party", SEMI_URL, SEMI_2)
inp("s_pkg_28", "SEMI 組裝與封裝設備 2028", "$B", None, 8.6, None, "Interested-party", SEMI_URL, SEMI_2 + "。2027 年值未找到 → 幾何內插（Derived）")
inp("s_tot_26", "SEMI 設備合計 2026（對帳用）", "$B", None, 165.9, None, "Interested-party", SEMI_URL, SEMI_2)
inp("g_tot_26", "SEMI 設備合計 2026 年增（對帳用）", "%", None, 0.232, None, "Interested-party", SEMI_URL, SEMI_2)
inp("s_tot_28", "SEMI 設備合計 2028（對帳用）", "$B", None, 229.5, None, "Interested-party", SEMI_URL, SEMI_2)
inp("g29", "2029 年 SEMI 測試、封裝設備年增（SEMI 無 2029 預測；僅用於測試與封裝）", "%", 0.0, 0.08, 0.12, "Assumed", "無。SEMI 前段年增 2026 +23.1%、2027 +21.8%、2028 +14.1%：逐年減速約 2–7 個百分點，外推 2029 約 +8%", "未找到第二來源；情境 A 的意思是「續增」，B、C 由情境係數覆寫")
sec("D. 資本支出情境係數（2028、2029；設備總額相對前一年的倍數；情境 A 沿用上表由下而上路徑，不用係數）")
inp("ck_B_28", "情境 B 2028 係數（持平）", "x", 0.95, 1.00, 1.05, "Assumed", "依核心問題 1：B＝2028 持平（循環在 2027 見頂）", "理由：記憶體模型 B 情境 2027 Q2 價格見頂；資本支出滯後約 1–2 季，2028 設備訂單轉為持平。無實證，非預測")
inp("ck_C_28", "情境 C 2028 係數（下滑約兩成）", "x", 0.75, 0.80, 0.85, "Assumed", "依核心問題 1：C＝2028 下滑約兩成", "理由：記憶體模型 C 情境 2028 價格回到 2025 水準；WFE 歷史下行年（2019、2023）跌幅較淺，C 為壓力情境。無實證")
inp("ck_B_29", "情境 B 2029 係數", "x", 0.90, 0.95, 1.05, "Assumed", "平台後的溫和消化", "無實證")
inp("ck_C_29", "情境 C 2029 係數", "x", 0.90, 1.00, 1.10, "Assumed", "谷底持平", "無實證")
sec("E. AI 相關分類（用於 AI 增量占比；v0.1 的分類假設，非公開資料）")
inp("adv_other", "台積電以外代工邏輯中屬先進邏輯的比例", "%", 0.15, 0.25, 0.40, "Assumed", "無：Intel 18A／14A、三星 SF2、Rapidus 相對成熟製程擴產（中芯、華虹等）的占比未找到", "未找到第二來源")
inp("pkg_adv", "組裝與封裝設備中屬先進封裝的比例", "%", 0.35, 0.50, 0.65, "Assumed", "無：KLA 稱 2026 先進封裝製程控制營收約 $1.1B，但整體先進封裝設備金額未找到", "未找到第二來源。台積電資本支出內的先進封裝設備可能與此重疊（見 README 限制）")
sec("F. 對帳來源值（本模型不得為對上這些數字而調整；2026 前段口徑各家不同，見 Recon）")
GS_URL = "https://www.gate.com/news/detail/goldman-sachs-significantly-raises-wfe-forecasts-through-2028-can-the-ai-23706576"
inp("rc_lam26", "Lam Research 2026 WFE（低 150B）", "$B", None, 150, None, "Interested-party", "Lam FQ4-26 法說逐字稿（https://www.investing.com/news/transcripts/earnings-call-transcript-lam-research-posts-record-q4-2026-results-stock-rebounds-93CH-4821978）；另：Citi 研討會 CFO 稱「low-$150B」（https://finance.yahoo.com/technology/ai/articles/lam-research-lifts-2026-wfe-180229282.html）", "Lam 2026-10-21 法說稿另稱「mid-$150B」（https://www.fool.com/earnings/call-transcripts/2026/10/21/lam-research-lrcx-q1-2027-earnings-call-transcript/）；取 150 為下限")
inp("rc_lam27", "Lam Research 2027 WFE（>200B，下限；額外對照）", "$B", None, 200, None, "Interested-party", "同上法說稿：2027 預估 >$200B", "使用者未列；額外列示")
inp("rc_kla26", "KLA 2026 WFE（低 150B；含先進封裝）", "$B", None, 150, None, "Interested-party", "KLA FQ4-26 法說摘要（社群摘要，弱）https://x.com/TheValueist/status/2082235400190640364；BigGo 摘要：https://finance.biggo.com/news/US_KLAC_2026-07-28", "社群摘要 https://x.com/TheValueist/status/2082235400190640364（弱）；KLA 2025 約 120 vs SEMI 116.9，口徑含先進封裝")
inp("rc_kla27", "KLA 2027 規劃（約 190B）", "$B", None, 190, None, "Interested-party", "KLA FQ4-26 法說摘要（社群摘要，弱）https://x.com/TheValueist/status/2082235400190640364", "同上社群摘要；一手逐字稿未找到")
inp("rc_ms27", "Morgan Stanley 2027 WFE", "$B", None, 202, None, "Interested-party", "二手轉述（itiger，2026-07-06：https://www.itiger.com/news/2649678863——「2027 由 191 上修到 202，2028 由 215 上修到 227」；本環境無法開啟該頁，未直接核對）", "上修前版本 191／215 與 Investing.com 報導一致（2026／27／28＝149／191／215；https://www.investing.com/news/stock-market-news/semicap-stocks-morgan-stanley-raises-wfe-outlook-and-shifts-ratings-4696580），故 202／227 為其上修值")
inp("rc_ms28", "Morgan Stanley 2028 WFE", "$B", None, 227, None, "Interested-party", "同上（itiger，二手轉述）", "同上")
inp("rc_gs26", "Goldman Sachs 2026 WFE", "$B", None, 150, None, "Interested-party", GS_URL, "https://www.kucoin.com/news/flash/goldman-sachs-upgrades-wfe-forecast-to-1-5-trillion-by-2026-dram-and-foundry-drive-growth（標題單位誤植，內文 150）")
inp("rc_gs27", "Goldman Sachs 2027 WFE", "$B", None, 218, None, "Interested-party", GS_URL, "同上")
inp("rc_gs28", "Goldman Sachs 2028 WFE", "$B", None, 281, None, "Interested-party", GS_URL, "同上")
inp("rc_thr", "差距門檻（README 須說明原因）", "%", None, 0.15, None, "Verified", "使用者指定 15%", "—")

R = {}; r = 4
for s in rows:
    if s[0] == "sec":
        IN.cell(r, 1, s[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = s
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    isdist = key.startswith("cap_") or key == "fx"
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            c = IN.cell(r, j, v); c.font = GREEN if isdist and j == 5 else BLUE
            if unit == "%": c.number_format = "0.0%"
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [14, 48, 9, 8, 10, 8, 18, 90, 90]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# ---------------- Mkt ----------------
M = wb.create_sheet("Mkt")
M["A1"] = "Mkt — 半導體設備市場（十億美元）：由下而上（公司資本支出 × 設備占比）＋ 其他廠（2025 基準 × 成長率）＋ SEMI 後段；三情境"
M["A2"] = "2025＝由下而上（三家＋台積電資本支出 × 設備占比）＋其他廠基準（SEMI 實績減已知廠）；2026 起其他廠＝基準 × 成長率（Assumed），總額由下而上，不由 SEMI 2026 以後預測決定。SEMI 2026 以後的分項只用於對帳。"
hdr(M, 3, ["項目", "單位"] + [str(2000 + y) for y in YRS] + ["說明"])
for y in YRS: M[f"{YC[y]}3"] = 2000 + y; M[f"{YC[y]}3"].number_format = "0"; M[f"{YC[y]}3"].font = BOLD; M[f"{YC[y]}3"].fill = HDR
MR = {}; mr = [40]   # 摘要區佔 4–38；明細自 40 起
def mput(label, unit, fd, key=None, fmt="0.0", note="", bold=False):
    rr = mr[0]
    M.cell(rr, 1, label).font = Font(bold=bold); M.cell(rr, 2, unit)
    for y, f in fd.items():
        if f is None: continue
        M[f"{YC[y]}{rr}"] = f; M[f"{YC[y]}{rr}"].number_format = fmt
    M.cell(rr, 8, note)
    if key: MR[key] = rr
    mr[0] += 1
    return rr
def msec(t): M.cell(mr[0], 1, t).font = BOLD; mr[0] += 1
c = lambda key, y: f"{YC[y]}{MR[key]}"
Y2 = (26, 27, 28, 29)

msec("1. SEMI 分項金額（A 路徑的天花板與「其他廠」推導基礎）")
def semi_row(k, label, g_chain=True):
    fd = {26: f"=s_{k}_26", 25: f"=D{{r}}/(1+g_{k}_26)", 27: f"=D{{r}}*(1+g_{k}_27)", 28: f"=E{{r}}*(1+g_{k}_28)", 29: f"=F{{r}}*(1+g29)"}
    rr = mr[0]; fd = {y: v.replace("{r}", str(rr)) for y, v in fd.items()}
    mput(label, "$B", fd, key=f"semi_{k}", note="2026 新聞稿值；2025 由年增回推；2027、2028 年增連乘；2029 為 g29（Assumed）")
for k, lb in (("dram", "SEMI DRAM 設備"), ("nand", "SEMI NAND 設備"), ("fd", "SEMI 代工與邏輯設備"), ("wfe", "SEMI 晶圓前段設備（WFE）")): semi_row(k, lb)
for k, lb in (("test", "SEMI 測試設備"), ("pkg", "SEMI 組裝與封裝設備")):
    rr = mr[0]
    mput(lb, "$B", {26: f"=s_{k}_26", 25: f"=D{rr}/(1+g_{k}_26)", 28: f"=s_{k}_28", 27: f"=SQRT(D{rr}*F{rr})", 29: f"=F{rr}*(1+g29)"}, key=f"semi_{k}",
         note="2026、2028 新聞稿值；2027 幾何內插（Derived，未找到年中預測 2027 值）")
mput("SEMI 其他前段（WFE − 代工邏輯 − DRAM − NAND）", "$B", {y: f"={c('semi_wfe', y)}-{c('semi_fd', y)}-{c('semi_dram', y)}-{c('semi_nand', y)}" for y in YRS}, key="semi_ofe", note="Derived：類比、功率、光罩、廠務設備等；v0.1 不由下而上，沿用 SEMI 差額")
mput("SEMI 設備合計（WFE＋測試＋封裝）", "$B", {y: f"={c('semi_wfe', y)}+{c('semi_test', y)}+{c('semi_pkg', y)}" for y in YRS}, key="semi_tot", bold=True, note="2026 應＝165.9、2028 應＝229.5（捨入差見 Recon）")
_r = mr[0]
mput("2025 年 SEMI DRAM ÷（DRAM＋NAND）（各年沿用，Assumed）", "%", {25: f"={c('semi_dram', 25)}/({c('semi_dram', 25)}+{c('semi_nand', 25)})", **{y: f"=$C{_r}" for y in Y2}}, key="dr", fmt="0.0%", note="把各廠記憶體設備拆成 DRAM（含 HBM）與 NAND：假設三家組合同 2025 年 SEMI 市場組合")
mput("SEMI 其他前段占（記憶體＋代工邏輯）比例（2025；各年沿用，Assumed）", "%", {25: f"={c('semi_ofe', 25)}/({c('semi_dram', 25)}+{c('semi_nand', 25)}+{c('semi_fd', 25)})", **{y: f"=$C{_r+1}" for y in Y2}}, key="ofer", fmt="0.0%", note="其他前段（類比、功率、光罩等）v0.1 維持 2025 年占比")
mr[0] += 1
msec("2. 記憶體晶圓前段設備（2025 以實績校準；2026 起沿用校準值與成長率）")
msec("2a. 2025 年其他記憶體廠：由下而上獨立估計（資本支出 × sh_oth）＝其他廠 2025 基準")
mput("長鑫資本支出", "$B", {25: "=ind_cxmt"}, key="i_cx", note="TechInsights（Reuters）；高 12.0 年度未核對，見 In／Sens")
mput("南亞科資本支出", "$B", {25: "=ind_nanya/fx_twd"}, key="i_ny")
mput("華邦資本支出", "$B", {25: 0}, key="i_wb", note="未找到，記 0")
mput("其他 DRAM 廠資本支出", "$B", {25: f"={c('i_cx', 25)}+{c('i_ny', 25)}+{c('i_wb', 25)}"}, key="i_dc")
mput("鎧俠資本支出（FY25）", "$B", {25: "=ind_kioxia/fx_jpy"}, key="i_kx")
mput("長江存儲資本支出（全球 NAND 資本支出 × 占比）", "$B", {25: "=ind_nand_ind*ind_ymtc_sh"}, key="i_ym")
mput("SanDisk（JV 分攤與鎧俠重疊，不另計）", "$B", {25: 0}, key="i_sd", note="FY25 未找到")
mput("其他 NAND 廠資本支出", "$B", {25: f"={c('i_kx', 25)}+{c('i_ym', 25)}+{c('i_sd', 25)}"}, key="i_nc")
mput("其他 DRAM 廠設備＝資本支出 × sh_oth", "$B", {25: f"={c('i_dc', 25)}*sh_oth"}, key="i_de")
mput("其他 NAND 廠設備＝資本支出 × sh_oth", "$B", {25: f"={c('i_nc', 25)}*sh_oth"}, key="i_ne")
mput("其他記憶體廠設備（獨立估計）", "$B", {25: f"={c('i_de', 25)}+{c('i_ne', 25)}"}, key="i_eq", bold=True)
mr[0] += 1
msec("2b. 三大廠資本支出")
mput("SK 海力士資本支出（美元）", "$B", {y: f"=cap_sk_{y}*1000/fx" for y in YRS}, key="sk", note="兆韓元 × 1000 ÷ KRW/USD ＝ 十億美元（2025 亦用同一匯率，Assumed）")
mput("美光資本支出（淨；FY≈日曆年）", "$B", {y: f"=cap_mu_{y}" for y in YRS}, key="mu")
mput("三星記憶體資本支出（美元）", "$B", {25: "=cap_ssds_25*ss_mem*1000/fx", **{y: f"=cap_ss_{y}*1000/fx" for y in Y2}}, key="ss", note="2025＝DS 47.5 兆 × 記憶體占比（Assumed）")
mput("三家資本支出合計", "$B", {y: f"={c('sk', y)}+{c('mu', y)}+{c('ss', y)}" for y in YRS}, key="cap3", bold=True)
mr[0] += 1
msec("2c. 校準（只用 2025 實績）：SK 海力士、三星共用一個設備占比，使 2025 年各廠設備加總＝SEMI 2025 DRAM＋NAND")
mput("SEMI 2025 DRAM＋NAND 設備（實績）", "$B", {25: f"={c('semi_dram', 25)}+{c('semi_nand', 25)}"}, key="cal_semi")
mput("  減：美光 2025 設備（資本支出 × sh_mu_25）", "$B", {25: f"={c('mu', 25)}*sh_mu_25"}, key="cal_mu")
mput("  減：其他記憶體廠 2025 設備（獨立估計）", "$B", {25: f"={c('i_eq', 25)}"}, key="cal_oth")
mput("  ＝SK 海力士＋三星 2025 設備（校準目標）", "$B", {25: f"={c('cal_semi', 25)}-{c('cal_mu', 25)}-{c('cal_oth', 25)}"}, key="cal_tgt")
mput("  ÷ SK 海力士＋三星 2025 資本支出", "$B", {25: f"={c('sk', 25)}+{c('ss', 25)}"}, key="cal_den")
mput("校準值：SK 海力士、三星設備占比（Derived；2026 年以後沿用）", "%", {25: f"={c('cal_tgt', 25)}/{c('cal_den', 25)}"}, key="cal", fmt="0.0%", bold=True, note="用公式算出；In 頁 sh_sk、sh_ss 基準取此值")
mput("  低：長鑫取高值（In 高）時的校準值", "%", {25: f"=({c('cal_semi', 25)}-({c('i_dc', 25)}-{c('i_cx', 25)}+ind_cxmt_hi+{c('i_nc', 25)})*sh_oth-{c('cal_mu', 25)})/{c('cal_den', 25)}"}, key="cal_lo", fmt="0.0%")
mput("  高：業界下限（sh_oth 低）", "%", {25: "=sh_oth_lo"}, key="cal_hi", fmt="0.0%")
mput("  對照：業界值（sh_oth 基準，不再當基準）", "%", {25: "=sh_oth"}, key="cal_ind", fmt="0.0%")
mr[0] += 1
msec("2d. 各廠設備與合計")
mput("SK 海力士設備＝資本支出 × sh_sk", "$B", {y: f"={c('sk', y)}*sh_sk" for y in YRS}, key="eq_sk")
mput("美光設備＝資本支出 × sh_mu（FY25／FY26／FY27–29）", "$B", {25: f"={c('mu', 25)}*sh_mu_25", 26: f"={c('mu', 26)}*sh_mu_26", **{y: f"={c('mu', y)}*sh_mu_27" for y in (27, 28, 29)}}, key="eq_mu")
mput("三星記憶體設備＝資本支出 × sh_ss", "$B", {y: f"={c('ss', y)}*sh_ss" for y in YRS}, key="eq_ss")
mput("三家設備合計", "$B", {y: f"={c('eq_sk', y)}+{c('eq_mu', y)}+{c('eq_ss', y)}" for y in YRS}, key="eq3", bold=True)
mput("  三家加權設備占比", "%", {y: f"={c('eq3', y)}/{c('cap3', y)}" for y in YRS}, key="wsh", fmt="0.0%")
mput("  其中 DRAM（× 2025 年 DRAM 占比）", "$B", {y: f"={c('eq3', y)}*{c('dr', y)}" for y in YRS}, key="eq3d", note="2025 年 DRAM／NAND 拆分為近似（以 SEMI 2025 比例拆三家），合計不受影響")
mput("  其中 NAND", "$B", {y: f"={c('eq3', y)}-{c('eq3d', y)}" for y in YRS}, key="eq3n")
_r = mr[0]
mput("其他 DRAM 廠設備（2025＝獨立估計；2026 起 × (1＋g_oth_dram)）", "$B", {25: f"={c('i_de', 25)}", **{y: f"={YC[y-1]}{_r}*(1+g_oth_dram)" for y in Y2}}, key="od")
_r = mr[0]
mput("其他 NAND 廠設備（2025＝獨立估計；2026 起 × (1＋g_oth_nand)）", "$B", {25: f"={c('i_ne', 25)}", **{y: f"={YC[y-1]}{_r}*(1+g_oth_nand)" for y in Y2}}, key="on")
mput("記憶體設備合計＝三家＋其他", "$B", {y: f"={c('eq3', y)}+{c('od', y)}+{c('on', y)}" for y in YRS}, key="mem", bold=True)
mput("核對：2025 記憶體合計 − SEMI 實績（應為 0）", "$B", {25: f"={c('mem', 25)}-{c('cal_semi', 25)}"}, key="chk_cal", fmt="0.000")
mput("診斷：記憶體合計 − SEMI DRAM＋NAND", "$B", {y: f"={c('mem', y)}-{c('semi_dram', y)}-{c('semi_nand', y)}" for y in YRS}, key="dmem", note="診斷值（不是輸入）：2026 起正＝本模型高於 SEMI；2025 為 0（校準）")
mr[0] += 1
msec("3. 代工與邏輯晶圓前段設備（2025：台積電由下而上＋SEMI 殘差求其他廠基準；2026 起：其他廠＝基準 × 成長率）")
mput("台積電資本支出", "$B", {y: f"=cap_ts_{y}" for y in YRS}, key="ts")
mput("台積電設備＝資本支出 × sh_fd", "$B", {y: f"={c('ts', y)}*sh_fd" for y in YRS}, key="eqts", note="v0.1.2：sh_fd 基準 64% 維持（Analogy），區間改 50–75%；證據與可識別性見 TSMC 頁。未另扣先進封裝／光罩（10–20%），上限影響見 Sens")
mput("2025 殘差：其他代工邏輯＝SEMI 代工邏輯 − 台積電設備（照實顯示）", "$B", {25: f"={c('semi_fd', 25)}-{c('eqts', 25)}"}, key="fo25", note="Derived：三星晶圓代工、Intel、中芯等。v0.1.2：此殘差無法拆成『台積電占比』與『未列名其他廠』（TSMC 頁可識別性檢查：需 X＝18–30 的未列名設備），所以台積電占比不由 SEMI 校準")
_r = mr[0]
mput("其他代工與邏輯設備（基準＝MAX(殘差, 0)；2026 起 × (1＋g_oth_fd)）", "$B", {25: f"=MAX({c('fo25', 25)},0)", **{y: f"={YC[y-1]}{_r}*(1+g_oth_fd)" for y in Y2}}, key="fo")
mput("代工與邏輯設備合計", "$B", {y: f"={c('eqts', y)}+{c('fo', y)}" for y in YRS}, key="fdt", bold=True)
mput("診斷：代工邏輯合計 − SEMI 代工邏輯", "$B", {y: f"={c('fdt', y)}-{c('semi_fd', y)}" for y in YRS}, key="dfd")
mr[0] += 1
msec("3b. 其他前段（類比、功率、光罩、廠務等）：維持 2025 年占（記憶體＋代工邏輯）的比例")
mput("其他前段設備", "$B", {25: f"={c('semi_ofe', 25)}", **{y: f"={c('ofer', y)}*({c('mem', y)}+{c('fdt', y)})" for y in Y2}}, key="ofe", note="2025＝SEMI 差額；2026 起＝比例 × 本模型（記憶體＋代工邏輯）")
mr[0] += 1
msec("3c. 診斷：三大廠隱含設備占比（2025）")
mput("三大廠隱含設備占比（獨立估計）＝(SEMI − 其他廠獨立估計) ÷ 三大廠資本支出", "%", {25: f"=({c('cal_semi', 25)}-{c('i_eq', 25)})/{c('cap3', 25)}"}, key="i_impl", fmt="0.0%", bold=True, note="＝三家加權設備占比（校準恆等）")
mput("  其中 SK 海力士、三星校準值", "%", {25: f"={c('cal', 25)}"}, key="i_cal", fmt="0.0%")
mput("  對照：業界值 sh_oth（Sens 保留 64% 對照列）", "%", {25: f"={c('cal_ind', 25)}"}, key="i_ind", fmt="0.0%")
mput("旗標：校準值低於業界下限 58%（1＝是）", "", {25: f"=IF({c('cal', 25)}<sh_oth_lo,1,0)"}, key="i_flag", fmt="0", note="低於下限＝SK 海力士、三星資本支出中建築與無塵室比重高於業界值（與美光「增量多數為建築」方向一致，非證明）")
mr[0] += 1
msec("4. 測試與封裝設備（v0.1 以 SEMI 金額為輸入，不做由下而上）")
M.cell(mr[0], 1, "為何不由下而上：需要「每顆加速器的測試時間（最終測試、燒機 burn-in、系統級測試 SLT 的小時數）」與「每條 CoWoS／HBM 堆疊產線的鍵合機台數」，兩者公開資料皆未找到（Hanmi 等僅有訂單，無台數與產能對照）；以 SEMI 為輸入、並標明為單一來源。"); mr[0] += 1
mput("測試設備", "$B", {y: f"={c('semi_test', y)}" for y in YRS}, key="test")
mput("組裝與封裝設備", "$B", {y: f"={c('semi_pkg', y)}" for y in YRS}, key="pkg")
mr[0] += 1

def scen(sn, title):
    msec(f"5{sn}. 情境 {sn}：{title}")
    L = {}
    def sl(key, label, fa, note="", bold=False):
        fd = {}
        for y in YRS:
            if sn == "A": f = fa(y)
            elif fa(y) is None: f = None
            elif y <= 27: f = f"={YC[y]}{MR['A_'+key]}"      # 2025–2027 與 A 相同
            elif y == 28: f = f"=E{mr[0]}*ck_{sn}_28"
            else: f = f"=F{mr[0]}*ck_{sn}_29"
            fd[y] = f
        mput(label, "$B", fd, key=f"{sn}_{key}", note=note, bold=bold)
    sl("L1", "先進邏輯（台積電設備＋其他代工中屬先進者）〔AI〕", lambda y: f"={c('eqts', y)}+adv_other*{c('fo', y)}")
    sl("L2", "其他代工與邏輯", lambda y: f"={c('fdt', y)}-{YC[y]}{mr[0]-1}")
    sl("L3", "DRAM（含 HBM）〔AI〕", lambda y: f"={c('eq3d', y)}+{c('od', y)}")
    sl("L4", "NAND", lambda y: f"={c('eq3n', y)}+{c('on', y)}")
    sl("L5", "其他前段", lambda y: f"={c('ofe', y)}")
    sl("L6", "先進封裝〔AI〕", lambda y: f"={c('pkg', y)}*pkg_adv")
    sl("L7", "測試與其他封裝", lambda y: f"={c('test', y)}+{c('pkg', y)}-{YC[y]}{mr[0]-1}")
    k = lambda x: f"{sn}_{x}"
    mput("晶圓前段設備", "$B", {y: f"=SUM({YC[y]}{MR[k('L1')]}:{YC[y]}{MR[k('L5')]})" for y in YRS}, key=k("wfe"))
    mput("後段設備（測試＋封裝）", "$B", {y: f"={YC[y]}{MR[k('L6')]}+{YC[y]}{MR[k('L7')]}" for y in YRS}, key=k("bk"))
    mput(f"設備總額（情境 {sn}）", "$B", {y: f"={c(k('wfe'), y)}+{c(k('bk'), y)}" for y in YRS}, key=k("tot"), bold=True)
    mput("  年增", "%", {y: f"={c(k('tot'), y)}/{c(k('tot'), y-1)}-1" for y in (26, 27, 28, 29)}, key=k("yoy"), fmt="0.0%", note="2026 年增含口徑轉換（2025＝SEMI 實績，2026＝由下而上）")
    mput("  前段年增", "%", {y: f"={c(k('wfe'), y)}/{c(k('wfe'), y-1)}-1" for y in (26, 27, 28, 29)}, key=k("wyoy"), fmt="0.0%")
    mput("AI 相關設備（先進邏輯＋DRAM／HBM＋先進封裝）", "$B", {y: f"={c(k('L1'), y)}+{c(k('L3'), y)}+{c(k('L6'), y)}" for y in YRS}, key=k("ai"), bold=True)
    mput("  AI 相關占總額", "%", {y: f"={c(k('ai'), y)}/{c(k('tot'), y)}" for y in YRS}, key=k("aish"), fmt="0.0%")
    mput("  AI 相關增量占總增量（相對 2025 累計）", "%", {y: f'=IF(ABS({c(k("tot"), y)}-{c(k("tot"), 25)})<0.05,"n/a",({c(k("ai"), y)}-{c(k("ai"), 25)})/({c(k("tot"), y)}-{c(k("tot"), 25)}))' for y in (26, 27, 28, 29)}, key=k("aid_cum"), fmt="0%", note="分子分母同為負（下滑）＝AI 相關占下滑額的比例；總增量接近 0 時比值失真（可超過 100%）")
    mput("  AI 相關增量占總增量（當年）", "%", {y: f'=IF(ABS({c(k("tot"), y)}-{c(k("tot"), y-1)})<0.05,"n/a",({c(k("ai"), y)}-{c(k("ai"), y-1)})/({c(k("tot"), y)}-{c(k("tot"), y-1)}))' for y in (26, 27, 28, 29)}, key=k("aid_yoy"), fmt="0%", note="增量為 0（持平）時不適用")
    tr = f"{YC[25]}{MR[k('tot')]}:{YC[29]}{MR[k('tot')]}"
    mx = f"MAX({tr})"
    ptot = lambda y: f"ROUND({YC[y]}{MR[k('tot')]},9)=ROUND({mx},9)"
    pk = "=" + "".join(f"IF({ptot(y)},{2000+y}," for y in (29, 28, 27, 26)) + "2025" + ")" * 4
    mput("見頂年度（設備總額最高水準；同值取較晚者）", "年", {25: pk}, key=k("peak"), fmt="0", note="若為 2029＝區間末尾，未見頂（'≥2029'）")
    mr[0] += 1
scen("A", "由下而上路徑（資本支出沿用記憶體與台積電模型；其他廠沿 SEMI；2029 沿 g29）")
scen("B", "2028 持平（In 係數 ck_B_28、ck_B_29 套用於 A 路徑的 2027 水準）")
scen("C", "2028 下滑約兩成（In 係數 ck_C_28、ck_C_29）")
msec("6. 核對")
mput("SEMI 設備合計 − 公布值（2026；捨入差）", "$B", {26: f"={c('semi_tot', 26)}-s_tot_26", 28: f"={c('semi_tot', 28)}-s_tot_28"}, key="chk", fmt="0.00", note="2028 差 −0.1 為 SEMI 分項捨入")
mput("2025 本模型（由下而上）− SEMI 實績合計", "$B", {25: f"={c('A_tot', 25)}-{c('semi_tot', 25)}"}, key="chk4", fmt="0.000", note="殘差為負時基準取 0，故可 ≠ 0")
mput("情境 A 分項合計 − 總額（應為 0）", "$B", {y: f"=SUM({YC[y]}{MR['A_L1']}:{YC[y]}{MR['A_L7']})-{c('A_tot', y)}" for y in Y2}, key="chk2", fmt="0.000")
mput("B、C 在 2025–2027 與 A 相同（差額合計，應為 0）", "$B", {y: f"=({c('B_tot', y)}-{c('A_tot', y)})+({c('C_tot', y)}-{c('A_tot', y)})" for y in (25, 26, 27)}, key="chk3", fmt="0.000")
last_row = mr[0]
# In 頁 sh_sk、sh_ss：基準／低／高取 Mkt 校準列
for r_ in range(4, IN.max_row + 1):
    if IN.cell(r_, 1).value in ("sh_sk", "sh_ss"):
        for col_, key_ in ((4, "cal_lo"), (5, "cal"), (6, "cal_hi")):
            cc = IN.cell(r_, col_, f"=Mkt!$C${MR[key_]}"); cc.number_format = "0.0%"; cc.font = Font(color="000000")

# 摘要區（4–38）
srow = 4
M.cell(srow, 1, "摘要（各情境設備總額、年增、AI 增量占比、見頂年度）").font = BOLD; srow += 1
SUMR = {}
for sn in "ABC":
    M.cell(srow, 1, f"情境 {sn}").font = BOLD; srow += 1
    for lab, key, fmt, yrs in (("設備總額", "tot", "0.0", YRS), ("  年增", "yoy", "0.0%", Y2), ("  其中晶圓前段", "wfe", "0.0", YRS), ("  AI 相關設備", "ai", "0.0", YRS),
                               ("  AI 增量占總增量（相對 2025 累計）", "aid_cum", "0%", (26, 27, 28, 29)), ("  AI 增量占總增量（當年）", "aid_yoy", "0%", (26, 27, 28, 29))):
        M.cell(srow, 1, lab); M.cell(srow, 2, "%" if fmt.endswith("%") else "$B")
        for y in yrs:
            M[f"{YC[y]}{srow}"] = f"={YC[y]}{MR[sn+'_'+key]}"; M[f"{YC[y]}{srow}"].number_format = fmt
        SUMR[(sn, key)] = srow; srow += 1
    M.cell(srow, 1, "  見頂年度"); M[f"C{srow}"] = f"=C{MR[sn+'_peak']}"; M[f"C{srow}"].number_format = "0"; SUMR[(sn, "peak")] = srow; srow += 1
assert srow <= 39, srow
M.column_dimensions["A"].width = 58; M.column_dimensions["B"].width = 6
for col in "CDEFG": M.column_dimensions[col].width = 10
M.column_dimensions["H"].width = 80
M.freeze_panes = "C4"

# ---------------- TSMC（v0.1.2：台積電設備占比證據） ----------------
TS_ = wb.create_sheet("TSMC")
TS_["A1"] = "TSMC — 台積電設備占資本支出：證據、可識別性、先進封裝／光罩重疊（v0.1.2）。金額：百萬美元（除另註）。藍字＝本頁輸入"
TS_["A2"] = ("年報／20-F「不動產、廠房及設備」附註的分類增添（房屋及建築、機器設備、待驗設備及未完工程）：未找到——本環境無法開啟年報與 SEC 原文，只能讀搜尋摘錄；該附註公司有揭露，屬檢索缺口，不是公司不揭露。"
             "以下改用董事會每季「資本預算」分項（核准額，不是實際支出）；各次分項來自二手摘要，三項加總與公司公告總額相符（核對欄）。")
hdr(TS_, 4, ["董事會日期", "機器設備：先進製程", "機器設備：先進封裝／成熟／特殊", "不動產與資本化租賃資產", "三項合計", "公司公告總額", "合計−公告", "設備占比", "來源1", "來源2／備註", "標記"])
EV = [  # 日期, 先進, 封裝等, 不動產, 公告總額, 來源1, 來源2/備註
 ("2025-05-13", 9590, 308, 5350, 15247, "Investing.com 6-K 摘要（2025-06-25）https://www.investing.com/news/sec-filings/tsmc-reports-board-share-changes-and-new-capital-spending-approvals-93CH-4110161", "總額：Focus Taiwan https://focustaiwan.tw/business/202505130022；Tom's Hardware 逐季 15.247"),
 ("2025-08-12", 7800, 2000, 10900, 20657, "Panabee 摘要（只到 0.1B）https://www.panabee.com/news/tsmc-commits-us-20-7-billion-to-bolster-advanced-tech-and-manufacturing-capacity", "總額：Tom's Hardware 逐季 20.657；公司 6-K 原文分項：未能開啟核對"),
 ("2026-02-10", 18921, 4670, 21372, 44962, "Investing.com 6-K 摘要 https://www.investing.com/news/sec-filings/tsmc-reports-board-shareholding-changes-and-major-capital-approvals-for-february-93CH-4579333", "總額：Taipei Times https://www.taipeitimes.com/News/biz/archives/2026/02/11/2003852112；TipRanks 同分項。公司新聞稿決議文只列用途與總額，分項見於 6-K 摘要"),
 ("2026-05-12", 21013, 0, 10271, 31284, "TipRanks 摘要 https://www.tipranks.com/news/company-announcements/tsmc-details-heavy-capex-bond-issuance-and-share-capital-reduction-for-may-2026", "總額：Focus Taiwan https://focustaiwan.tw/business/202605120026；摘要只列兩項，封裝等機器設備未單獨列（可能併入先進製程）"),
 ("2026-08", 16035, 4791, 8616, 29442.5, "Stocktitan 6-K 摘要 https://www.stocktitan.net/sec-filings/TSM/6-k-taiwan-semiconductor-manufacturing-co-ltd-current-report-foreign--efeb868a1401.html", "總額：Focus Taiwan https://focustaiwan.tw/business/202608110015（29.44B）；SEC 檔 tsm-monthend6kx20260924（搜尋結果列出，未開啟）")]
for i, (d, a, b, rl, tot_, s1, s2) in enumerate(EV):
    rr = 5 + i; TS_.cell(rr, 1, d)
    for j, v in ((2, a), (3, b), (4, rl), (6, tot_)): TS_.cell(rr, j, v).font = BLUE; TS_.cell(rr, j).number_format = "#,##0.0"
    TS_.cell(rr, 5, f"=SUM(B{rr}:D{rr})").number_format = "#,##0.0"; TS_.cell(rr, 7, f"=E{rr}-F{rr}").number_format = "#,##0.0"
    TS_.cell(rr, 8, f"=(B{rr}+C{rr})/E{rr}").number_format = "0.0%"
    TS_.cell(rr, 9, s1); TS_.cell(rr, 10, s2); TS_.cell(rr, 11, "Interested-party（台積電揭露；二手摘要，合計與公告總額核對）")
for rr, lab, a, b in ((10, "2025 兩次合計（5、8 月）", 5, 6), (11, "2026 三次合計", 7, 9), (12, "五次合計", 5, 9)):
    TS_.cell(rr, 1, lab).font = BOLD
    for j in range(2, 7):
        col = "ABCDEF"[j - 1]; TS_.cell(rr, j, f"=SUM({col}{a}:{col}{b})").number_format = "#,##0.0"
    TS_.cell(rr, 8, f"=(B{rr}+C{rr})/E{rr}").number_format = "0.0%"; TS_.cell(rr, 11, "Derived")
TS_.cell(13, 1, "單次最低／最高設備占比"); TS_.cell(13, 8, "=MIN(H5:H9)").number_format = "0.0%"; TS_.cell(13, 9, "=MAX(H5:H9)").number_format = "0.0%"
TS_.cell(14, 1, ("怎麼讀：(1) 這是核准額，不是支出；核准在前、多年執行，且「資本化租賃資產」含非現金項目（2025 年四次核准合計 68.0B，現金資本支出 40.9B）。"
                 "(2) 建廠先於裝機，擴廠加速期的不動產占比會領先實際支出占比。(3) SEMI 的 WFE 含「廠務設備」（Fab Facilities），台積電『不動產與廠務系統』裡的廠務機電部分在 SEMI 口徑可能算設備，所以和 SEMI 可比的設備占比 ≥ 本欄。"
                 "(4) 2025 年只有 2 次、2026 年 3 次找到分項；2025-02、2025-11 與 2021–2024 各次：只找到總額（下表），分項未找到。"))
TS_.cell(14, 1).alignment = WRAP
hdr(TS_, 16, ["只找到總額的董事會核准（分項未找到）", "核准總額 US$M", "", "", "", "", "", "", "來源", "備註", "標記"])
TOT_ONLY = [("2025-02-12", 17141.4, "TSMC 決議新聞稿 https://pr.tsmc.com/english/news/3207", ""),
            ("2025-11-11", 14981, "Tom's Hardware 逐季；Tribune https://www.tribuneindia.com/news/business/tsmc-approved-a-capital-appropriation-plan-of-usd-14-98-bn-company-also-raised-dividend-payout/amp", "此次含 2026 年研發與經常性資本預算、資本化租賃資產"),
            ("2024-02", 9421.48, "搜尋摘要（TSMC 決議）", "含資本化租賃資產"), ("2024-06-05", 17356.2, "TSMC 決議 https://pr.tsmc.com/system/files/newspdf/attachment/4f8848e007a8291ca9e5e720193efebc599d0a9e/0605%20board%20meeting%20(E)_final_wmn.pdf", ""),
            ("2024-08-13", 29615.47, "MarketScreener 6-K 標題頁", ""), ("2024-11-12", 15479.95, "SEC 6-K https://www.sec.gov/Archives/edgar/data/1046179/000104617924000124/tsm-boardx20241112x6k.htm", ""),
            ("2023-02-14", 6959.5, "SEC 6-K https://www.sec.gov/Archives/edgar/data/1046179/000156459023001822/tsm-6k_20230214.htm", ""), ("2023-05-09", 366.1, "搜尋摘要（SEC 6-K）", "僅廠房建設"),
            ("2023-08-08", 6059.5, "SEC 6-K https://www.sec.gov/Archives/edgar/data/1046179/000162828023027944/tsm-boardx20230808x6k.htm", "2023-11 一次：未找到"),
            ("2022 全年", 20944.17, "TSMC 2022 年報 https://investor.tsmc.com/sites/ir/annual-report/2022/2022%20Annual%20Report-E.pdf（搜尋摘錄）", "2022-12 摘錄顯示「不動產與資本化租賃資產」曾單列約 1.1B，其餘分項未找到"),
            ("2021-06", 9290, "Focus Taiwan https://focustaiwan.tw/business/202106100005", "2021 其他季：未找到")]
for i, (d, v, s1, nt) in enumerate(TOT_ONLY):
    rr = 17 + i; TS_.cell(rr, 1, d); TS_.cell(rr, 2, v).font = BLUE; TS_.cell(rr, 2).number_format = "#,##0.0"; TS_.cell(rr, 9, s1); TS_.cell(rr, 10, nt); TS_.cell(rr, 11, "Verified（公司公告總額；搜尋摘錄）")
rr = 17 + len(TOT_ONLY) + 1
TS_.cell(rr, 1, "現金資本支出（取得不動產、廠房及設備）").font = BOLD; hdr(TS_, rr + 1, ["年度", "新台幣百萬元", "美元百萬元", "", "", "", "", "", "來源", "備註", "標記"])
CASH = [(2021, None, None, "未找到（本次搜尋未取得 2021 年數字）", ""), (2022, 1082672, None, "TSMC 2024 20-F 搜尋摘錄（2022–2024 資本支出）https://investor.tsmc.com/sites/ir/sec-filings/2024%2020-F.pdf", ""),
        (2023, 949817, None, "同上", ""), (2024, 956007, 29755, "同上；4Q24 Management Report 29.76B", "美元為公司換算"),
        (2025, 1272410.529, 40900, "Stocktitan 2025 年報 6-K 摘要（PP&E 增添 1,272,410,529 千元）；美元 40.9B：Semiconductor Intelligence https://www.semiconductorintelligence.com/capex-up-for-foundry-memory/", "新台幣為二手摘要，未開原文")]
for i, (y, n, u, s1, nt) in enumerate(CASH):
    r2 = rr + 2 + i; TS_.cell(r2, 1, y)
    for j, v in ((2, n), (3, u)):
        if v is not None: TS_.cell(r2, j, v).font = BLUE; TS_.cell(r2, j).number_format = "#,##0"
    TS_.cell(r2, 9, s1); TS_.cell(r2, 10, nt); TS_.cell(r2, 11, "Verified（現金流量表金額）" if n else "未找到")
rr = rr + 2 + len(CASH) + 1
TS_.cell(rr, 1, "其他證據與業界值").font = BOLD; hdr(TS_, rr + 1, ["項目", "低", "基準／中", "高", "", "", "", "", "來源", "備註", "標記"])
OTH = [("元大投顧：先進製程外商前段設備占資本支出", 0.70, None, 0.80, "自由財經／財富自由（2026-01 法說後）https://stock.ltn.com.tw/article/70szg3dsjge2", "媒體轉述，原報告全文未找到；只含外商主導的前段設備、限先進製程；投顧報告以設備供應鏈受惠股為主題，立場偏多", "Interested-party（元大投顧）"),
       ("業界值：SEMI 2025 WFE ÷ 全球資本支出", 0.58, 0.64, 0.70, "見 In 頁 SH_SRC1／SH_SRC2（SEMI；Semiconductor Intelligence；Rebound Capital）", "業界合計，非台積電", "Analogy（全產業，差異：含記憶體與 IDM）"),
       ("SK 海力士、三星校準值（對照）", 0.385, "=sh_sk", 0.58, "Mkt 2c", "Derived，記憶體廠；用來問「台積電是否也同樣偏低」", "Derived")]
OTH_R0 = rr + 2
for i, (nm, lo, mid, hi, s1, nt, tg) in enumerate(OTH):
    r2 = rr + 2 + i; TS_.cell(r2, 1, nm)
    for j, v in ((2, lo), (3, mid), (4, hi)):
        if v is not None: TS_.cell(r2, j, v).number_format = "0.0%"; TS_.cell(r2, j).font = BLUE if not isinstance(v, str) else Font(color="000000")
    TS_.cell(r2, 9, s1); TS_.cell(r2, 10, nt); TS_.cell(r2, 11, tg)
rr = rr + 2 + len(OTH) + 1
TS_.cell(rr, 1, "可識別性：能不能用 SEMI 2025 代工邏輯實績校準台積電的設備占比？（比照 SK＋三星，但其他廠無法獨立估計）").font = BOLD; rr += 1
ID0 = rr
rows_id = [("SEMI 2025 代工與邏輯設備（實績）", f"=Mkt!C{MR['semi_fd']}", "$B", "In 頁／Mkt 1"),
           ("台積電 2025 資本支出", "=cap_ts_25", "$B", "In"),
           ("已知其他廠 2025 資本支出：中芯＋Intel（毛額）", "=cap_smic_25+cap_intel_25", "$B", "僅此二家有數字；三星晶圓代工（報導約 5 兆韓元，未核對）、聯電（資料站約 1.5B，弱）、GlobalFoundries、華虹、力積電、世界先進、Rapidus、中國未上市廠：未找到可用的 2025 資本支出"),
           ("已知其他廠設備（× 業界值 64%）", f"=B{ID0+2}*sh_oth", "$B", "Analogy：假設同業界值；Intel 含大量建廠，可能偏高")]
for i, (a, f, u, nt) in enumerate(rows_id):
    TS_.cell(rr, 1, a); TS_.cell(rr, 2, f).number_format = "0.0"; TS_.cell(rr, 3, u); TS_.cell(rr, 10, nt); rr += 1
hdr(TS_, rr, ["假設台積電設備占比 s", "台積電設備", "需要的『未列名其他廠設備』X＝SEMI−台積電−已知其他廠", "X 占 SEMI 代工邏輯", "", "", "", "", "", "說明", ""]); rr += 1
for s_ in ("sh_fd_lo", "sh_fd", "sh_fd_hi", "sh_sk"):
    TS_.cell(rr, 1, f"={s_}").number_format = "0.0%"; TS_.cell(rr, 2, f"=B{ID0+1}*A{rr}").number_format = "0.0"
    TS_.cell(rr, 3, f"=B{ID0}-B{rr}-B{ID0+3}").number_format = "0.0"; TS_.cell(rr, 4, f"=C{rr}/B{ID0}").number_format = "0%"; rr += 1
TS_.cell(rr, 1, "s 的上限情況：X＝0 時隱含的 s"); TS_.cell(rr, 2, f"=(B{ID0}-B{ID0+3})/B{ID0+1}").number_format = "0%"
TS_.cell(rr, 10, "超過 100% 不可能：SEMI 代工邏輯設備出貨，大於「台積電全部資本支出＋已知其他廠 × 64%」。"); rr += 1
TS_.cell(rr, 1, ("結論：不論 s 取 45.5%、64% 或 75%，都需要 X＝約 18–30 十億美元（占 SEMI 代工邏輯 28–46%）的『未列名其他廠設備』才能湊成 SEMI 實績，而這個 X 沒有任何獨立估計（未找到）。"
                 "因此與 SK＋三星不同：記憶體校準有 SEMI 實績＋美光（一手揭露）＋其他廠獨立估計三個已知量，只剩一個未知數；代工邏輯是『台積電 s』與『X』兩個未知數、一條方程式，不可識別。"
                 "SEMI 實績只約束 s×台積電資本支出＋其他廠＝65.6，s 由其他證據決定。可能的來源：SEMI 計設備出貨、資本支出是現金支出（時差，見限制 5）；中國未上市晶圓廠與口徑差異；本表已知其他廠的資本支出不全。"))
TS_.cell(rr, 1).alignment = WRAP; rr += 2
TS_.cell(rr, 1, "先進封裝、測試、光罩與 SEMI 分類的重疊").font = BOLD; rr += 1
OV = [("公司稱用於先進封裝、測試、光罩及其他（pk_ts）", "=pk_ts_lo", "=pk_ts", "=pk_ts_hi", "%"),
      ("2026 台積電資本支出（Mkt，三情境相同）", f"=Mkt!D{MR['ts']}", f"=Mkt!D{MR['ts']}", f"=Mkt!D{MR['ts']}", "$B"),
      ("該類資本支出 × 設備占比 sh_fd（若全部被 SEMI 分入後段或其他前段）", f"=B{rr+1}*B{rr}*sh_fd", f"=C{rr+1}*C{rr}*sh_fd", f"=D{rr+1}*D{rr}*sh_fd", "$B"),
      ("占情境 A 2026 設備總額", f"=B{rr+2}/Mkt!D{MR['A_tot']}", f"=C{rr+2}/Mkt!D{MR['A_tot']}", f"=D{rr+2}/Mkt!D{MR['A_tot']}", "%")]
for a, f1, f2, f3, u in OV:
    TS_.cell(rr, 1, a)
    for j, f in ((2, f1), (3, f2), (4, f3)): TS_.cell(rr, j, f).number_format = "0.0%" if u == "%" else "0.0"
    rr += 1
TS_.cell(rr, 1, ("判斷：SEMI 的口徑——總設備＝WFE（晶圓處理、光罩、廠務）＋測試＋組裝封裝（來源：SEMI 簡報定義，搜尋摘錄）；台積電的先進封裝（CoWoS 等）設備應落在 SEMI『組裝與封裝』或『測試』，光罩設備落在 WFE 的光罩項（本模型的『其他前段』），"
                 "而本模型把台積電全部資本支出 × 設備占比放進『代工邏輯』，另外再以 SEMI 金額加入測試與封裝——所以有重疊的可能，上表是上限（10–20% 內含測試、光罩與『其他』，且部分封裝用的蝕刻、沉積、微影設備 SEMI 可能算在 WFE）。"
                 "2025 年因代工邏輯由 SEMI 殘差補平，重疊不影響總額；2026 年以後台積電成長（+52%）遠高於其他廠（+8%），重疊使總額高估；直接扣除為 2026 年設備總額的 2.2–4.5%（上表），但因 2025 殘差補平，Sens 的淨影響較小：情境 A 2026–2029 約 −1.1% 至 −2.2%（扣 15%／20%）。SEMI 是否把 CoWoS 前段工序算在 WFE：未找到。base 不扣，理由：重疊程度無法識別，且扣除方向會被『未列名其他廠』的校準殘差吸收。"))
TS_.cell(rr, 1).alignment = WRAP
for col, w in zip("ABCDEFGHIJK", [46, 16, 22, 20, 12, 12, 10, 10, 70, 70, 34]): TS_.column_dimensions[col].width = w

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型與外部來源對帳。不為對上來源而調整參數；差距 > 15% 的原因見 README。模型欄＝晶圓前段設備（除非另註）；來源口徑各異（SEMI WFE、分析師 WFE、KLA 含先進封裝）"
hdr(RC, 3, ["項目", "年度", "來源", "來源值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "A 差距 > 門檻", "口徑備註"])
def scn(sn, key, y): return f"=Mkt!{YC[y]}{MR[sn+'_'+key]}"
items = []  # (label, year, source, source_expr, key(modelrow), yfn?, note)
def add(label, y, src, expr, mk, note="", growth=False):
    items.append((label, y, src, expr, mk, note, growth))
add("設備合計", 2026, "SEMI 2026-07", "=s_tot_26", "tot", "SEMI 公布 165.9", )
add("晶圓前段", 2026, "SEMI 2026-07", "=s_wfe_26", "wfe", "")
add("代工與邏輯", 2026, "SEMI 2026-07", "=s_fd_26", ("fd",), "模型＝代工邏輯設備合計（台積電由下而上＋其他殘差）")
add("DRAM", 2026, "SEMI 2026-07", "=s_dram_26", ("dram",), "模型記憶體設備 × SEMI DRAM 占比")
add("NAND", 2026, "SEMI 2026-07", "=s_nand_26", ("nand",), "")
add("晶圓前段年增", 2027, "SEMI 2026-07（+21.8%）", "=g_wfe_27", "wfeyoy", "年增率之差以百分點計", True)
add("晶圓前段", 2027, "SEMI 2026-07（由 2026 值連乘年增）", f"=Mkt!E{MR['semi_wfe']}", "wfe", "SEMI 2027 絕對金額為推得（143.9×1.218）")
add("晶圓前段", 2028, "SEMI 2026-07（連乘）", f"=Mkt!F{MR['semi_wfe']}", "wfe", "新聞稿稱「達 200B」")
add("設備合計", 2028, "SEMI 2026-07", "=s_tot_28", "tot", "公布 229.5；分項加總 229.4")
add("晶圓前段", 2026, "Lam Research（低 150B）", "=rc_lam26", "wfe", "Lam 2026-10 法說另稱 mid-150")
add("晶圓前段", 2026, "KLA（低 150B；含先進封裝）", "=rc_kla26", "wfe", "KLA 口徑含先進封裝，2025 基準約 120 vs SEMI 116.9")
add("晶圓前段", 2027, "KLA 規劃（約 190B）", "=rc_kla27", "wfe", "")
add("晶圓前段", 2027, "Lam Research（>200B；額外對照）", "=rc_lam27", "wfe", "取下限")
add("晶圓前段", 2027, "Morgan Stanley（itiger 二手轉述 https://www.itiger.com/news/2649678863）", "=rc_ms27", "wfe", "itiger 二手轉述（2026-07-06）：2027 由 191 上修到 202；連結見 In 頁")
add("晶圓前段", 2028, "Morgan Stanley（itiger 二手轉述 https://www.itiger.com/news/2649678863）", "=rc_ms28", "wfe", "itiger 二手轉述：2028 由 215 上修到 227")
add("晶圓前段", 2026, "Goldman Sachs", "=rc_gs26", "wfe", "")
add("晶圓前段", 2027, "Goldman Sachs", "=rc_gs27", "wfe", "")
add("晶圓前段", 2028, "Goldman Sachs", "=rc_gs28", "wfe", "")
rr = 4
for label, y, src, expr, mk, note, growth in items:
    RC.cell(rr, 1, label); RC.cell(rr, 2, y); RC.cell(rr, 3, src); RC.cell(rr, 4, expr)
    for j, sn in zip((5, 7, 9), "ABC"):
        if mk == "wfeyoy": f = f"=Mkt!{YC[y-2000]}{MR[sn+'_wyoy']}"
        elif mk == "tot" or mk == "wfe": f = f"=Mkt!{YC[y-2000]}{MR[sn+'_'+mk]}"
        elif mk == ("fd",): f = f"=Mkt!{YC[y-2000]}{MR[sn+'_L1']}+Mkt!{YC[y-2000]}{MR[sn+'_L2']}"
        elif mk == ("dram",): f = f"=Mkt!{YC[y-2000]}{MR[sn+'_L3']}"
        elif mk == ("nand",): f = f"=Mkt!{YC[y-2000]}{MR[sn+'_L4']}"
        RC.cell(rr, j, f)
        RC.cell(rr, j + 1, f"={'E' if j==5 else 'G' if j==7 else 'I'}{rr}-D{rr}" if growth else f"={'E' if j==5 else 'G' if j==7 else 'I'}{rr}/D{rr}-1")
        RC.cell(rr, j + 1).number_format = "0.0%"
    fmt = "0.0%" if growth else "0.0"
    for j in (4, 5, 7, 9): RC.cell(rr, j).number_format = fmt
    RC.cell(rr, 11, f'=IF(ABS(F{rr})>rc_thr,"超過門檻","")' if not growth else f'=IF(ABS(F{rr})>0.05,"超過 5 個百分點","")')
    RC.cell(rr, 12, note); rr += 1
RC.cell(rr + 1, 1, "核對：SEMI 2026 分項加總").font = BOLD
RC.cell(rr + 2, 1, "DRAM＋NAND＋代工邏輯＋其他前段 − WFE"); RC.cell(rr + 2, 4, f"=Mkt!D{MR['semi_dram']}+Mkt!D{MR['semi_nand']}+Mkt!D{MR['semi_fd']}+Mkt!D{MR['semi_ofe']}-Mkt!D{MR['semi_wfe']}").number_format = "0.000"
RC.cell(rr + 3, 1, "注意：校準只用 2025 實績；2026 年以後 A 路徑的總額由下而上（其他廠＝2025 基準 × 成長率，Assumed），不由 SEMI 或券商預測決定（測試與封裝除外，v0.1 以 SEMI 為輸入），所以對 SEMI 與各來源的差距有意義。")
for col, w in zip("ABCDEFGHIJKL", [14, 7, 36, 10, 10, 9, 10, 9, 10, 9, 16, 70]): RC.column_dimensions[col].width = w
RECON_LAST = rr

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（低／高取 In 的區間）；輸出：情境 A 2027、2028、2029 設備總額（$B）。最後兩組列出「其他記憶體廠 2025 基準」若不為 0 對總額的影響（2025 殘差為負時基準取 0）"
hdr(SE, 3, ["輸入", "2027 A", "相對基準", "2028 A", "相對基準", "2029 A", "相對基準", "說明"])
def parts(**p):
    P = dict(sk="sh_sk", ss="sh_ss", mu25="sh_mu_25", mu26="sh_mu_26", mu27="sh_mu_27", fd="sh_fd", fx="fx", gd="g_oth_dram", gn="g_oth_nand", gf="g_oth_fd",
             sk25="cap_sk_25", ssm="ss_mem", ts25="cap_ts_25", cx="ind_cxmt", sho="sh_oth", pk="0", recal=False)
    P.update(p)
    if P["recal"]:   # 2025 輸入改變時重新校準 SK、三星共用占比
        oth = f"((Mkt!$C${MR['i_dc']}-Mkt!$C${MR['i_cx']}+{P['cx']}+Mkt!$C${MR['i_nc']})*{P['sho']})"
        mu = f"(cap_mu_25*{P['mu25']})"
        den = f"({P['sk25']}*1000/{P['fx']}+cap_ssds_25*{P['ssm']}*1000/{P['fx']})"
        P["sk"] = P["ss"] = f"((Mkt!$C${MR['cal_semi']}-{oth}-{mu})/{den})"
    return P
def eq3(y, P):
    mu_sh = {25: P["mu25"], 26: P["mu26"]}.get(y, P["mu27"])
    sk = f"{P['sk25'] if y == 25 else 'cap_sk_'+str(y)}*1000/{P['fx']}*{P['sk']}"
    mu = f"cap_mu_{y}*{mu_sh}"
    ss = (f"cap_ssds_25*{P['ssm']}*1000/{P['fx']}*{P['ss']}" if y == 25 else f"cap_ss_{y}*1000/{P['fx']}*{P['ss']}")
    return f"({sk}+{mu}+{ss})"
def tot(y, P):
    base_d = f"((Mkt!$C${MR['i_dc']}-Mkt!$C${MR['i_cx']}+{P['cx']})*{P['sho']})"
    base_n = f"(Mkt!$C${MR['i_nc']}*{P['sho']})"
    n = y - 25
    mem = f"({eq3(y, P)}+({base_d})*(1+{P['gd']})^{n}+({base_n})*(1+{P['gn']})^{n})"
    ts = f"{P['ts25'] if y == 25 else 'cap_ts_'+str(y)}*(1-{P['pk']})*{P['fd']}"
    fd = f"({ts}+MAX(Mkt!$C${MR['semi_fd']}-{P['ts25']}*(1-{P['pk']})*{P['fd']},0)*(1+{P['gf']})^{n})"
    return f"={mem}+{fd}+Mkt!{YC[y]}{MR['ofer']}*({mem}+{fd})+Mkt!{YC[y]}{MR['test']}+Mkt!{YC[y]}{MR['pkg']}"
cases = [("基準（2025 實績校準）", {}),
 ("SK 海力士＋三星設備占比 低（長鑫取 12.0 的校準值）", {"sk": "sh_sk_lo", "ss": "sh_ss_lo", "cx": "ind_cxmt_hi"}),
 ("SK 海力士＋三星設備占比 高（58%，業界下限）", {"sk": "sh_sk_hi", "ss": "sh_ss_hi"}),
 ("對照：SK 海力士＋三星設備占比＝業界值 64%（不再當基準）", {"sk": "sh_oth", "ss": "sh_oth"}),
 ("長鑫 2025 資本支出取高 12.0（年度未核對；重新校準）", {"cx": "ind_cxmt_hi", "recal": True}),
 ("美光設備占比 低（FY25／26／27–29；重新校準）", {"mu25": "sh_mu_25_lo", "mu26": "sh_mu_26_lo", "mu27": "sh_mu_27_lo", "recal": True}),
 ("美光設備占比 高（重新校準）", {"mu25": "sh_mu_25_hi", "mu26": "sh_mu_26_hi", "mu27": "sh_mu_27_hi", "recal": True}),
 ("其他記憶體廠設備占比 58%（重新校準）", {"sho": "sh_oth_lo", "recal": True}), ("其他記憶體廠設備占比 70%（重新校準）", {"sho": "sh_oth_hi", "recal": True}),
 ("台積電設備占比 低（50%）", {"fd": "sh_fd_lo"}), ("台積電設備占比 高（75%）", {"fd": "sh_fd_hi"}),
 ("台積電設備占比＝SK 海力士、三星校準值（45.5%；若台積電也同樣偏低）", {"fd": "sh_sk"}),
 ("台積電資本支出扣除先進封裝／光罩等 10%（設備同比例；上限，見 TSMC 頁）", {"pk": "pk_ts_lo"}), ("台積電資本支出扣除先進封裝／光罩等 20%（上限）", {"pk": "pk_ts_hi"}),
 ("匯率 低（1,350；重新校準）", {"fx": "fx_lo", "recal": True}), ("匯率 高（1,500；重新校準）", {"fx": "fx_hi", "recal": True}),
 ("SK 海力士 2025 資本支出 低（28 兆；重新校準）", {"sk25": "cap_sk_25_lo", "recal": True}),
 ("三星 DS 記憶體占比 65%（重新校準）", {"ssm": "ss_mem_lo", "recal": True}), ("三星 DS 記憶體占比 90%（重新校準）", {"ssm": "ss_mem_hi", "recal": True}),
 ("台積電 2025 資本支出 低（40）", {"ts25": "cap_ts_25_lo"}), ("台積電 2025 資本支出 高（42）", {"ts25": "cap_ts_25_hi"}),
 ("其他 DRAM 廠成長 低", {"gd": "g_oth_dram_lo"}), ("其他 DRAM 廠成長 高", {"gd": "g_oth_dram_hi"}),
 ("其他 NAND 廠成長 低", {"gn": "g_oth_nand_lo"}), ("其他 NAND 廠成長 高", {"gn": "g_oth_nand_hi"}),
 ("其他代工邏輯成長 低", {"gf": "g_oth_fd_lo"}), ("其他代工邏輯成長 高", {"gf": "g_oth_fd_hi"})]
for i, (nm, kw) in enumerate(cases):
    rr = 4 + i; SE.cell(rr, 1, nm)
    for j, y in ((2, 27), (4, 28), (6, 29)):
        SE.cell(rr, j, tot(y, parts(**kw))).number_format = "0.0"
        SE.cell(rr, j + 1, f"={'BDF'[(j-2)//2]}{rr}/{'BDF'[(j-2)//2]}$4-1").number_format = "0.0%"
SE.cell(4, 8, "同 Mkt 情境 A（公式重算；核對列應為 0）")
SE.cell(4 + 14, 8, "標「重新校準」的列：2025 輸入改變後重新解出 SK、三星共用占比。匯率改變時校準值同比例改變，設備金額不變（匯率對 2026 年以後設備無影響）")
rr = 4 + len(cases)
SE.cell(rr, 1, "核對：基準 − Mkt 情境 A 總額（應為 0）")
for j, y in ((2, 27), (4, 28), (6, 29)): SE.cell(rr, j, f"={'BDF'[(j-2)//2]}4-Mkt!{YC[y]}{MR['A_tot']}").number_format = "0.000"
rr += 2
hdr(SE, rr, ["情境係數", "2028 設備總額", "相對 2027", "2029 設備總額", "相對 2028", "", "", "說明"]); rr += 1
for nm, sn, k28, k29 in (("B 係數 低", "B", "ck_B_28_lo", "ck_B_29_lo"), ("B 係數 基準", "B", "ck_B_28", "ck_B_29"), ("B 係數 高", "B", "ck_B_28_hi", "ck_B_29_hi"),
                         ("C 係數 低", "C", "ck_C_28_lo", "ck_C_29_lo"), ("C 係數 基準", "C", "ck_C_28", "ck_C_29"), ("C 係數 高", "C", "ck_C_28_hi", "ck_C_29_hi")):
    SE.cell(rr, 1, nm); SE.cell(rr, 2, f"=Mkt!E{MR[sn+'_tot']}*{k28}").number_format = "0.0"; SE.cell(rr, 3, f"=B{rr}/Mkt!E{MR[sn+'_tot']}-1").number_format = "0.0%"
    SE.cell(rr, 4, f"=B{rr}*{k29}").number_format = "0.0"; SE.cell(rr, 5, f"=D{rr}/B{rr}-1").number_format = "0.0%"; rr += 1
SE.cell(rr, 1, "說明：係數為 Assumed，區間無實證；B 的見頂年度對 2029 係數敏感（≥1.00 則 2029 不下滑，見頂年度移到 2029）。")
rr += 2
SE.cell(rr, 1, "校準值 vs 業界值 64%（對照）：三情境總額與見頂年度").font = BOLD; rr += 1
hdr(SE, rr, ["情境／占比", "2025", "2026", "2027", "2028", "2029", "見頂年度", "說明"]); rr += 1
def peak(row):
    mx = f"MAX(B{row}:F{row})"
    return "=" + "".join(f"IF(ROUND({'BCDEF'[k]}{row},9)=ROUND({mx},9),{2025+k}," for k in (4, 3, 2, 1)) + "2025" + ")" * 4
blkrows = {}
for lab, P in (("校準值（基準）", parts()), ("業界值 64%", parts(sk="sh_oth", ss="sh_oth"))):
    for sn in "ABC":
        SE.cell(rr, 1, f"情境 {sn}（{lab}）")
        for k, y in enumerate((25, 26, 27, 28, 29)):
            col = "BCDEF"[k]
            if sn == "A" or y <= 27: SE[f"{col}{rr}"] = tot(y, P)
            elif y == 28: SE[f"{col}{rr}"] = f"=D{rr}*ck_{sn}_28"
            else: SE[f"{col}{rr}"] = f"=E{rr}*ck_{sn}_29"
            SE[f"{col}{rr}"].number_format = "0.0"
        SE[f"G{rr}"] = peak(rr); SE[f"G{rr}"].number_format = "0"
        blkrows[(lab, sn)] = rr; rr += 1
for sn in "ABC":
    SE.cell(rr, 1, f"  業界值 64% − 校準值：情境 {sn}")
    for col in "BCDEF": SE[f"{col}{rr}"] = f"={col}{blkrows[('業界值 64%', sn)]}/{col}{blkrows[('校準值（基準）', sn)]}-1"; SE[f"{col}{rr}"].number_format = "0.0%"
    rr += 1
rr += 2
SE.cell(rr, 1, "台積電設備占比／封裝光罩扣除（v0.1.2）：三情境總額與見頂年度（其他參數不動）").font = BOLD; rr += 1
hdr(SE, rr, ["情境／台積電設定", "2025", "2026", "2027", "2028", "2029", "見頂年度", "說明"]); rr += 1
ts_blk = {}
for lab, P in (("基準 64%", parts()), ("50%（低）", parts(fd="sh_fd_lo")), ("75%（高）", parts(fd="sh_fd_hi")), ("45.5%（比照 SK、三星）", parts(fd="sh_sk")), ("基準 64% 並扣封裝／光罩 15%", parts(pk="pk_ts")), ("基準 64% 並扣 20%（上限）", parts(pk="pk_ts_hi"))):
    for sn in "ABC":
        SE.cell(rr, 1, f"情境 {sn}（台積電 {lab}）")
        for k, y in enumerate((25, 26, 27, 28, 29)):
            col = "BCDEF"[k]
            if sn == "A" or y <= 27: SE[f"{col}{rr}"] = tot(y, P)
            elif y == 28: SE[f"{col}{rr}"] = f"=D{rr}*ck_{sn}_28"
            else: SE[f"{col}{rr}"] = f"=E{rr}*ck_{sn}_29"
            SE[f"{col}{rr}"].number_format = "0.0"
        SE[f"G{rr}"] = peak(rr); SE[f"G{rr}"].number_format = "0"
        ts_blk[(lab, sn)] = rr; rr += 1
for lab in ("50%（低）", "75%（高）", "45.5%（比照 SK、三星）", "基準 64% 並扣封裝／光罩 15%", "基準 64% 並扣 20%（上限）"):
    for sn in "ABC":
        SE.cell(rr, 1, f"  相對基準 64%：情境 {sn}（{lab}）")
        for col in "BCDEF": SE[f"{col}{rr}"] = f"={col}{ts_blk[(lab, sn)]}/{col}{ts_blk[('基準 64%', sn)]}-1"; SE[f"{col}{rr}"].number_format = "0.0%"
        rr += 1
SE.cell(rr, 1, "說明：2025 因代工邏輯由 SEMI 殘差補平，台積電占比不影響 2025 總額；2026 年以後台積電資本支出成長（+52%）遠高於其他廠（+8%，Assumed），占比降低使總額下降。見頂年度由情境係數決定，不受台積電占比影響（三情境見頂年度與基準相同，見 G 欄）。")
rr += 2
SE.cell(rr, 1, "說明：業界值 64% 下 2025 本模型總額高於 SEMI 實績（三家 64% 加上其他廠獨立估計超過 SEMI）；校準值使 2025 恰等於 SEMI 實績。單因子列中未標「重新校準」者，校準值固定不變。")
SE.column_dimensions["A"].width = 48
for col in "BCDEFG": SE.column_dimensions[col].width = 12
SE.column_dimensions["H"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源"; SO["A1"].font = BOLD
hdr(SO, 2, ["來源", "網址／位置", "用途", "證據強度"])
srcs = [
 ("SEMI 年中設備預測新聞稿 2026-07-14", "https://www.semi.org/en/semi-press-release/global-semiconductor-equipment-sales-forecast-to-reach-a-record-229-billion-dollars-in-2028-semi-reports", "2026–2028 設備金額、年增", "中（一手但協會有立場；分項年增取自搜尋摘要，未能直接開啟全文）"),
 ("SEMI 新聞稿轉載（PR Newswire）", "https://www.prnewswire.co.uk/news-releases/global-semiconductor-equipment-sales-forecast-to-reach-a-record-229-billion-in-2028-semi-reports-302824214.html", "同上第二來源（非獨立）", "中"),
 ("Semiconductor Digest 轉載", "https://www.semiconductor-digest.com/global-semiconductor-equipment-sales-forecast-to-reach-a-record-229-billion-in-2028-semi-reports/", "同上", "中"),
 ("Semiconductor Intelligence：CapEx Up for Foundry, Memory", "https://www.semiconductorintelligence.com/capex-up-for-foundry-memory/", "2025 全球半導體資本支出 166B、台積電 40.9B", "中"),
 ("Rebound Capital：WFE 產業深度", "https://reboundcapital.substack.com/p/industry-deep-dive-semiconductor", "資本支出→WFE 轉換率約 58%", "弱（Substack）"),
 ("SK 海力士、美光、三星、台積電模型 dist 成品", "skhynix／micron／samsung／tsmc 的 dist/*.xlsx（Y 或 PL 頁「資本支出」列）", "公司資本支出", "各模型 In 頁標記"),
 ("Qz：SK 海力士 2Q26（2025 資本支出 30.2 兆）", "https://qz.com/sk-hynix-q2-2026-earnings-record-profit-misses-estimates-072926", "SK 海力士 2025 資本支出", "弱～中（摘要）"),
 ("MLQ：SK 海力士 2Q26", "https://mlq.ai/news/sk-hynix-posts-record-605-trillion-won-operating-profit-raises-capacity-spending/", "同上第二來源", "弱～中"),
 ("三星 IR：FY2025 CAPEX Results", "https://www.samsung.com/global/ir/reports-disclosures/public-disclosure-view.84547/", "DS 2025 資本支出 47.5 兆", "強（公司一手；頁面日期與內文不一致，需核對）"),
 ("美光 FY25 8-K／10-K", "https://www.sec.gov/Archives/edgar/data/723125/000072312525000024/a2025q4ex991-pressrelease.htm", "FY25 淨資本支出 13.80", "強"),
 ("TrendForce：鎧俠 FY26–28 資本支出", "https://www.trendforce.com/news/2026/06/03/news-kioxia-targets-%C2%A5470b-yearly-capex-in-fy26-28-up-66-from-fy25-reportedly-weighing-third-kitakami-fab-and-ma/", "其他 NAND 廠成長率", "中"),
 ("Motley Fool：SanDisk／鎧俠日本投資", "https://www.fool.com/investing/2026/08/29/sandisk-and-kioxia-plan-to-invest-more-than-usd31-billion-in-japanese-memory-plants-about-60-of-what-they-have-spent-there-in-25-years/", "SanDisk 毛資本支出約營收 6%", "弱～中"),
 ("Digitimes：南亞 2027 資本支出", "https://www.digitimes.com/news/a20260713PD211/nanya-technology-capex-capacity-2027.html", "其他 DRAM 廠成長率", "中"),
 ("TrendForce：中芯 2026 資本支出持平", "https://www.trendforce.com/news/2026/03/30/news-smic-2026-action-plan-points-to-above-industry-growth-capex-held-in-line-with-2025/", "其他代工邏輯成長率", "中"),
 ("Gate：Intel 2026 資本支出 >$20B", "https://www.gate.com/news/detail/intel-raises-2026-capex-guidance-to-over-20-billion-q2-earnings-beat-22949346", "其他代工邏輯成長率", "弱～中"),
 ("Micron FQ3 2026 準備稿", "https://s25.q4cdn.com/621799436/files/doc_financials/2026/q3/Q3-FY26-Prepared-Remarks.pdf", "FY27 資本支出增量多數為建築（無塵室）", "強（公司一手）"),
 ("Lam Research 2026-10-21 法說稿", "https://www.fool.com/earnings/call-transcripts/2026/10/21/lam-research-lrcx-q1-2027-earnings-call-transcript/", "2026 mid-150、2027 >200", "中"),
 ("Lam 提高 2026 WFE 預期（Yahoo）", "https://finance.yahoo.com/technology/ai/articles/lam-research-lifts-2026-wfe-180229282.html", "2026 low-150", "中"),
 ("KLA FQ4 2026 法說摘要（BigGo）", "https://finance.biggo.com/news/US_KLAC_2026-07-28", "2026 low-150、2027 約 190", "弱～中（轉述）"),
 ("KLA 法說摘要（X）", "https://x.com/TheValueist/status/2082235400190640364", "同上；交期 18–24 個月", "弱（社群貼文）"),
 ("Morgan Stanley WFE 上修（Investing.com，較早版本）", "https://www.investing.com/news/stock-market-news/semicap-stocks-morgan-stanley-raises-wfe-outlook-and-shifts-ratings-4696580", "149／191／215（舊版）；202／227 未能查到", "中"),
 ("Goldman Sachs WFE 上修（Gate）", "https://www.gate.com/news/detail/goldman-sachs-significantly-raises-wfe-forecasts-through-2028-can-the-ai-23706576", "150／218／281", "弱～中（轉述）"),
 ("TrendForce 2026-09-18 設備零組件交期", "https://www.trendforce.com/news/2026/09/18/news-chip-equipment-component-lead-times-reportedly-double-stretching-to-40-months-as-price-pressures-mount/", "零組件交期加倍、最長 40 個月", "中（'reportedly'）"),
 ("BigGo：日本關鍵零件交期 40 個月", "https://finance.biggo.com/news/f46a8b91-046a-4708-87db-63b34e0862b2", "同上第二來源", "弱～中"),
 ("ASML 2026 Q2 6-K", "https://www.sec.gov/Archives/edgar/data/0000937966/000162828026048235/pressreleasefinancialresul.htm", "EUV 產能、訂單", "強（公司一手）"),
 ("ASML Q3 2026：EUV 需求超過 2027 產能（Pomegra）", "https://pomegra.io/news/asml-q3-2026-euv-demand-outruns-2027-capacity", "第二來源", "弱"),
 ("TrendForce：SK 海力士 HBM4 TC 鍵合機訂單", "https://www.trendforce.com/news/2025/12/12/news-sk-hynix-reportedly-places-hbm4-tc-bonder-order-with-asmpt-as-hanmi-hanwha-patent-clash-heats-up/", "鍵合機供應商集中度", "中"),
 ("Futurum：KLA FQ2 2026（客戶受廠房就緒度限制）", "https://futurumgroup.com/insights/kla-corporation-q2-fy-2026-results-beat-on-ai-led-demand/", "無塵室／廠房限制", "中"),
 ("Samsung 投資 P5 無塵室（Cleanroom Technology）", "https://cleanroomtechnology.com/samsung-speeds-up-construction-of-p5-cleanroom-for", "無塵室先行", "中"),
 ("台積電董事會資本預算（Investing.com／TipRanks／Stocktitan／Panabee 對 6-K 的摘要）", "見 TSMC 頁來源欄", "台積電設備占比：機器設備 vs 不動產與資本化租賃資產（v0.1.2）", "中（公司揭露，二手摘要；合計與公司公告總額核對相符；6-K 原文未開啟）"),
 ("元大投顧：先進製程外商前段設備占資本支出 7～8 成（自由財經）", "https://stock.ltn.com.tw/article/70szg3dsjge2", "台積電設備占比第二來源", "弱～中（媒體轉述，投顧立場偏多，原報告未找到）"),
 ("台積電 3Q25 法說逐字稿（Motley Fool）", "https://www.fool.com/earnings/call-transcripts/2025/10/17/tsmc-tsm-q3-2025-earnings-call-transcript", "資本支出配置：先進製程約 70%、特殊 10–20%、先進封裝／測試／光罩 10–20%", "中（公司一手，經逐字稿網站）"),
 ("SEMI 設備口徑簡報（搜尋摘錄）", "https://www.semi.org/sites/semi.org/files/2024-07/SEMI%20Market%20Outlook_2024%20SEMICON%20WEST_SMG%20Meetingl.pdf", "總設備＝WFE（晶圓處理＋光罩＋廠務）＋測試＋組裝封裝", "中（摘錄，未開啟）"),
 ("Intel 4Q25 法說摘要（Yahoo Finance）", "https://finance.yahoo.com/news/intel-q4-earnings-call-highlights-234233695.html", "Intel 2025 資本支出毛額 17.7、扣抵 6.5（TSMC 頁可識別性檢查）", "中（摘要）"),
]
for i, row in enumerate(srcs, 3):
    for j, v in enumerate(row, 1): SO.cell(i, j, v)
for col, w in zip("ABCD", [52, 120, 48, 30]): SO.column_dimensions[col].width = w

# ---------------- README ----------------
BOTTLE = [
 ("無塵室空間", "美光：FY27 資本支出『增量』一半以上為建築（FQ3 稿）、『增量』多數為建築且用於 2028 下半年後的無塵室（FQ4 法說）；KLA 稱客戶受廠房就緒度與廠房外殼限制，壓住 1H 出貨；SK 證券稱記憶體供給成長受無塵室限制至 2027；SK 海力士首座無塵室提前到 2027-02", "Micron FQ3／FQ4 FY26 稿；Futurum（KLA FQ2 2026 摘要）；Cleanroom Technology（三星 P5）", "強（美光一手）／中／弱", "解釋為何三家資本支出 × 設備占比可能超過 SEMI 記憶體總額；設備認列落後資本支出現金支出一至兩季，v0.1 未建模"),
 ("長交期零組件", "TrendForce：設備零組件（閥件管路、陶瓷、RF 電源、氣體箱）交期據報倍增、最長 40 個月；BigGo：日本關鍵零件最長 40 個月，沉積設備零組件由約 4 個月增至約 10 個月", "TrendForce 2026-09-18；BigGo Finance", "中（皆為「據報」）／弱；應用材料、Lam 的一手說法：未找到", "限制 2027–28 設備出貨上限，不改變需求面；若上限生效，情境 A 的 2028 可能偏高"),
 ("EUV", "ASML：2026 低 NA 產能約 65 台、2027 再增約 30%；2027、2028 需求超過供給（二手摘要）；各來源對 2027 台數不一致（56／80／約 85）", "ASML 2026 Q2 6-K；Pomegra、TechPowerUp、Tom's 類二手整理", "強（產能，一手）／弱（2027 台數與供不應求，二手）", "EUV 台數限制先進邏輯與 DRAM 的 2027–28 增量；v0.1 未建模"),
 ("量測", "KLA 稱 2026 先進封裝製程控制營收約 $1.1B；交期 18–24 個月僅見於社群貼文轉述 KLA 法說；「量測產能是瓶頸」本身的證據：未找到（搜尋 Onto、KLA 皆無交期或積壓一手數字）", "Zacks／TradingView；X（TheValueist）；Futurum", "弱；瓶頸證據：未找到", "不影響 v0.1 數字；列為待查"),
 ("HBM 鍵合機", "Hanmi 占 HBM TC 鍵合機營收 71.2%（TechInsights，2025 前三季）；SK 海力士 HBM4 約用 50 組 TC 鍵合機、約半數來自 ASMPT；預期約 100 台訂單（M15X 1b DRAM）；Hanmi 6 月接 SK 海力士約 15 台。供給短缺或交期：未找到", "TrendForce 2025-12-12；The Elec；Techtimes 2026-06", "中（供應商集中度）；短缺：未找到", "鍵合機屬後段／封裝，v0.1 以 SEMI 金額為輸入，未建模"),
]
RECON_TXT = [
 "對帳差距 > 15% 的項目與可能原因（v0.1 建置時數值；2025 實績校準之外，2026 年以後完全沒有參照 SEMI 或券商的預測來設定參數）：",
 "(a) 2026 NAND +25.9%（模型 17.5 vs SEMI 13.9）：2025 年三家設備以 SEMI 2025 的 DRAM／NAND 比例（72.4／27.6）拆分，其他 NAND 廠（鎧俠、長江存儲）2025 獨立估計 3.9 $B 已占 SEMI 2025 NAND 10.6 的 37%，再乘 18% 成長；SEMI 2026 NAND 占比降到 26.4%。DRAM 僅 +1.5%，DRAM＋NAND 合計高 4.2 $B（+7.9%）。2025 年 DRAM／NAND 拆分本身就不準（模型 DRAM 26.4 vs SEMI 27.9、NAND 12.2 vs 10.6），合計則因校準而相等。",
 "(b) Goldman Sachs 2028 −22.9%（模型 A 216.6 vs 281）：GS 2027 年增約 +45%（150→218）、2028 +29%，分別高於 SEMI 的 +21.8%、+14.1%；GS 為最高的一組，可能納入不同口徑（KLA 口徑含先進封裝；廠務設備）、更激進的 DRAM／先進代工假設，且發布較晚（2026-08，SEMI 為 2026-07）。",
 "(c) 未超過 15% 的差距：本模型高於 SEMI（2026 前段 +8.5%、2027 +10.0%、2028 +8.3%），低於 Lam／Morgan Stanley／GS 的 2027–2028 數字（−3.6% 到 −11.6%），落在 KLA 190（+1.5%）附近。B、C 在 2028 的差距是情境設計的結果。",
 "(d) 校準：SK 海力士、三星共用一個設備占比（45.5%，區間 38.5–58%），使 2025 年各廠設備加總等於 SEMI 2025 DRAM＋NAND 實績。校準值低於業界下限 58%，表示 SK 海力士、三星的資本支出中建築與無塵室比重高於業界值，與美光「增量多數為建築」的方向一致（方向，不是證明）。校準的前提是其他記憶體廠 2025 獨立估計（8.6 $B，長鑫取 12.0 時校準值降到 38.5%）與三星記憶體資本支出假設（DS 47.5 兆 × 75%，區間 65–90%）。代工邏輯未比照辦理（台積電沒有類似美光的揭露、2025 殘差為正，維持 64%）；代工邏輯是否也有建築比重偏高的問題，待查。",
]
rr = RECON_LAST + 5
RC.cell(rr, 1, "台積電設備占比對帳（v0.1.2；不併入上方差距表，單位：占比，差距為百分點）").font = BOLD
hdr(RC, rr + 1, ["項目", "", "來源", "來源值", "模型 sh_fd 基準", "差距（模型−來源）", "", "", "", "", "", "口徑備註"])
for i, (lab, src, ref, nt) in enumerate((
        ("董事會資本預算：五次合計設備占比", "台積電 6-K 二手摘要（TSMC 頁）", "=TSMC!H12", "核准額，非支出；含資本化租賃資產；單次 47–71%"),
        ("董事會資本預算：2025 兩次（5、8 月）", "同上", "=TSMC!H10", "只有 2 次找到分項；2025 全年支出占比：未找到"),
        ("董事會資本預算：2026 三次（2、5、8 月）", "同上", "=TSMC!H11", ""),
        ("元大投顧：先進製程外商前段設備（中點）", "自由財經轉述（Interested-party）", f"=AVERAGE(TSMC!B{OTH_R0},TSMC!D{OTH_R0})", "只含外商前段設備、限先進製程"),
        ("業界值（SEMI WFE ÷ 全球資本支出等）", "In 頁 sh_oth", "=sh_oth", "Analogy，模型基準的來源"),
        ("SK 海力士、三星校準值（對照，不是台積電的證據）", "Mkt 2c", "=sh_sk", "記憶體廠，Derived")), 1):
    r2 = rr + 1 + i
    RC.cell(r2, 1, lab); RC.cell(r2, 3, src); RC.cell(r2, 4, ref).number_format = "0.0%"; RC.cell(r2, 5, "=sh_fd").number_format = "0.0%"
    RC.cell(r2, 6, f"=E{r2}-D{r2}").number_format = "0.0%"; RC.cell(r2, 12, nt)
RD = wb.create_sheet("README", 0)
RD_LINES = [
 ("半導體設備市場模型 v0.1.2（AI 半導體 Project 步驟 5：設備與材料——設備部分）　2026-10-23", True),
 ("模型命題：設備需求＝為滿足記憶體、代工擴產所需的資本支出 × 設備占比；本模型回答核心問題 1——設備循環在 2027 還是 2028 見頂——並列出核心問題 2、3 的瓶頸表。公司份額與毛利留給公司層，v0.1 不做。", False), ("", False),
 ("驅動因子 → 推導量", True),
 ("In：公司資本支出（讀自 skhynix／micron／samsung／tsmc dist，不修改它們）、設備占資本支出比例（兩個來源給低／高）、SEMI 2026-07 分項、情境係數（Assumed）、AI 分類假設", False),
 ("Mkt：2025 年 SK 海力士、三星設備占比以 SEMI 實績校準（Derived），其他記憶體廠 2025 基準＝獨立估計，代工邏輯其他廠＝SEMI 實績殘差；2026 起各廠資本支出 × 占比，其他廠基準 × 成長率（Assumed）；測試與封裝＝SEMI；三情境 A／B／C；AI 增量占比；見頂年度", False),
 ("Recon：對 SEMI、Lam、KLA、Morgan Stanley、Goldman Sachs 的差距；不為對上而調參。Sens：占比、匯率、2025 資本支出、其他廠成長率與基準、情境係數", False), ("", False),
 ("結果摘要（v0.1 建置時；見 Mkt 摘要區）", True),
 ("設備總額（$B）：A 2025–2029＝134.7／178.2／218.2／246.0／260.7（2029 仍增，未在區間內見頂）；B＝134.7／178.2／218.2／218.2／207.3（見頂 2028，與 2027 同水準）；C＝134.7／178.2／218.2／174.6／174.6（見頂 2027）。", False),
 ("校準（只用 2025 實績）：SK 海力士、三星共用一個設備占比 45.5%（Derived；低 38.5%＝長鑫取 12.0，高 58%＝業界下限），使 2025 年「三家設備＋其他記憶體廠獨立估計」＝SEMI 2025 DRAM＋NAND 實績（38.5）；2026 年以後沿用。美光維持由原文推導的占比（FY26 52.1%、FY27–29 46.0%）；台積電維持 64%（Analogy）。業界值 64% 不再當基準，Sens 保留對照列。", False),
 ("2026 年以後的前段設備完全沒有參照 SEMI 或券商預測：三家與台積電的資本支出 × 設備占比，加「其他廠＝2025 獨立估計（記憶體）或 SEMI 實績殘差（代工邏輯）× 成長率（Assumed）」；其他前段維持 2025 年占比。測試與封裝為 SEMI 輸入（v0.1 規格）。SEMI 2026 以後的 DRAM、NAND、代工邏輯、前段預測只用於對帳。", False),
 ("AI 相關增量（先進邏輯＋DRAM／HBM＋先進封裝）占總增量（相對 2025 累計）：A 2026–2029＝63／63／61／60%；B 63／63／63／64%；C 63／63／72／72%。結論取決於情境係數（Assumed），不是預測。", False), ("", False),
 ("限制", True),
 ("(1) 校準依賴其他記憶體廠 2025 獨立估計（長鑫 6.9–12.0 無一手資料、長江存儲為推算、華邦與 SanDisk 未找到）與三星記憶體資本支出假設；校準值低於業界下限 58%，方向與美光「增量多數為建築」一致，但不是證明。2025 年 DRAM／NAND 拆分不準（合計準）。", False),
 ("(2) 其他廠成長率為 Assumed（長鑫、長江存儲、三星晶圓代工、華虹的資本支出指引未找到）。(3) 測試與封裝以 SEMI 為輸入：缺少每顆加速器的測試時間與每條產線鍵合機台數資料。(4) 情境係數為 Assumed，是「若 B／C 則…」的情境。", False),
 ("(5) Morgan Stanley 202／227 為 itiger 二手轉述（2026-07-06，環境無法開啟，未直接核對；上修前 191／215 另有 Investing.com 佐證）。(6) 台積電資本支出含先進封裝、光罩（公司稱 10–20%），與 SEMI 後段、其他前段可能重疊：重疊程度無法識別，base 不扣，上限影響見 Sens（情境 A 2026–2029 約 −1.1% 至 −2.2%）。", False),
 ("(7) 時差：資本支出是現金支出，設備營收是出貨認列，兩者約差一到兩季，v0.1 沒有建模（年度金額直接相乘）。(8) 設備占比：美光依公司揭露推導（引文為一手，比例為 Derived；公司原文是「增量」多數為建築，並未給設備占比）；台積電的設備／建築比例未找到年報分類數字（附註有揭露，本環境無法開啟原文），用業界值 64%（Analogy，區間改 50–75%），以董事會資本預算分項（5 次彙總 60%）與元大 7～8 成交叉檢查。", False), ("", False),
 ("v0.1.2（2026-10-23）台積電設備占比查證", True),
 ("結論：台積電設備占比維持 64%（Analogy），只改區間（58–70% → 50–75%）；數值不變，三情境設備總額與見頂年度與 v0.1.1 完全相同。理由：(a) 董事會資本預算「機器設備」占比 5 次彙總 60%（單次 47–71%），元大稱先進製程前段設備 7～8 成，兩者夾住 64%，都沒有像 SK 海力士、三星（45.5%）那樣明顯偏低；(b) SEMI 2025 代工邏輯實績無法校準台積電：台積電占比與『未列名其他廠』是兩個未知數、一條方程式（TSMC 頁），所以沒有比照記憶體做校準。", False),
 ("證據缺口：年報／20-F 不動產廠房及設備附註的分類增添（2021–2025）、董事會預算 2021–2024 與 2025-02、2025-11 的分項、2021 年現金資本支出：未找到（檢索缺口，非公司不揭露）。董事會預算是核准額不是支出，且含非現金的資本化租賃資產。", False), ("", False),
 ("瓶頸表（核心問題 2、3；v0.1 只做表，不入模）", True)]
for t in RECON_TXT: pass
r0 = len(RD_LINES) + 1
for i, (t, b) in enumerate(RD_LINES, 1):
    RD.cell(i, 1, t).font = Font(bold=b); RD.cell(i, 1).alignment = WRAP
for j, h in enumerate(["瓶頸", "現況（v0.1 能找到的證據）", "來源", "證據強度", "對本模型的含意"], 1):
    c_ = RD.cell(r0, j, h); c_.font = BOLD; c_.fill = HDR
for i, row in enumerate(BOTTLE, r0 + 1):
    for j, v in enumerate(row, 1): RD.cell(i, j, v).alignment = WRAP
rr = r0 + len(BOTTLE) + 2
RD.cell(rr, 1, "對帳差距說明").font = BOLD
for i, t in enumerate(RECON_TXT, rr + 1): RD.cell(i, 1, t).alignment = WRAP
rr += len(RECON_TXT) + 2
RD.cell(rr, 1, "標記：Verified、Interested-party、Analogy、Assumed、Derived；Analogy 與 Assumed 給區間。SemiAnalysis 未使用。本檔不含任何個人投資部位或個人先驗。")
for col, w in zip("ABCDE", [18, 80, 40, 30, 60]): RD.column_dimensions[col].width = w
if __name__ == "__main__" and "--md" in sys.argv:
    pass
wb.save(sys.argv[1]); print("saved", {k: v for k, v in MR.items() if k.endswith("_tot") or k.endswith("_peak")})
