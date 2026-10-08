# AI 網通市場模型 v0.1（AI 半導體 Project 步驟 3）
# 用法：python3 markets/network/build_network.py raw.xlsx ；再以 LibreOffice headless 重算
# 加速器顆數取自 markets/memory/market.py（與 HBM 市場同一份數量層）；每顆網通內容以 markets/chip_bridge v0.1 為錨
import sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
import market as MEM

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR

COMP = [("su_sw", "scale-up 交換晶片"), ("su_ic", "scale-up 互連與交換匣其他（銅背板、線纜、電路板、機構）"), ("so_sw", "scale-out 交換器"),
        ("so_opt", "scale-out 光模組"), ("so_cu", "scale-out 銅纜（DAC／AEC）"), ("nic", "網卡／DPU（SuperNIC）")]
# 每顆加速器網通內容（美元）：su_sw, su_ic, so_sw, so_opt, so_cu, nic
PLAT = {
 "GB":   ("GB300 NVL72（Blackwell）", (1737, 1737, 3125, 2430, 1389, 2624), "Derived", "橋接表 v0.1 GB300：NVLink 交換晶片 0.846、交換匣其他 0.846、機架外 scale-out 3.382、機架內網卡＋DPU 1.278（$B/GW）÷ 每 GW 487,008 顆；機架外依 交換器 45%／光模組 35%／銅 20% 拆——v0.1 初版 35／50／15 使 2026 光學 41.8B，高於 LightCounting 26B 61%，改以 LightCounting 與光模組廠營收加總（約 26–30B）校準 [Derived]。合計約 $13.0K；Bernstein GB200 約 $10.7K（第二來源）"),
 "VR":   ("VR200 NVL72（Rubin）", (2917, 2917, 5251, 4084, 2334, 7000), "Derived", "橋接表 v0.1 VR200 ÷ 每 GW 291,744 顆；NVLink 6 交換晶片 36 顆（GB300 的 2 倍）、ConnectX-9 每架 144 顆；機架外拆分同上（45／35／20）。合計約 $24.5K"),
 "RU":   ("Rubin Ultra（MGX 單架／NVL576）", (4375, 4375, 5251, 4084, 2334, 7000), "Assumed", "scale-up 為 VR200 的 1.5 倍（NVL576 兩層 NVLink）；scale-out 同 VR200"),
 "HP":   ("Hopper HGX（8 GPU）", (300, 100, 2520, 1960, 1120, 2000), "Assumed", "NVSwitch 機內；400G InfiniBand 每 GPU 1 埠；合計約 $8K"),
 "AMD8": ("AMD MI355（8 GPU 伺服器）", (300, 100, 2520, 1960, 1120, 2000), "Analogy", "比照 Hopper HGX"),
 "AMDR": ("AMD Helios（MI455X 機架）", (2334, 2334, 4201, 3267, 1867, 5600), "Analogy", "VR200 的 0.8 倍（UALink／乙太網路 scale-up）"),
 "TPU":  ("Google TPU", (800, 1200, 800, 1500, 300, 400), "Assumed", "ICI 互連＋光學電路交換（OCS）；合計約 $5K"),
 "TRN":  ("AWS Trainium", (600, 800, 800, 1200, 300, 300), "Assumed", "NeuronLink＋EFA；合計約 $4K"),
 "ASIC": ("其他 ASIC", (600, 800, 800, 1200, 300, 300), "Assumed", "比照 Trainium"),
}
ITU = {"GB": (76.9, "Derived", "Tokenomics v5.27 IF_CapexIT 37.445 $B ÷ 487,008"), "VR": (128.8, "Derived", "37.586 ÷ 291,744"),
       "RU": (257.2, "Derived", "50.086 ÷ 194,760"), "HP": (38.3, "Derived", "27.510 ÷ 718,048"), "AMD8": (40, "Assumed", ""),
       "AMDR": (100, "Assumed", "VR200 的約 0.8 倍"), "TPU": (25, "Assumed", ""), "TRN": (20, "Assumed", ""), "ASIC": (20, "Assumed", "")}
TYPE2PLAT = {"bw": {y: "GB" for y in (26, 27, 28, 29)}, "rb": {y: "VR" for y in (26, 27, 28, 29)}, "ru": {y: "RU" for y in (26, 27, 28, 29)},
             "hp": {y: "HP" for y in (26, 27, 28, 29)}, "amd": {26: "AMD8", 27: "AMDR", 28: "AMDR", 29: "AMDR"},
             "tpu": {y: "TPU" for y in (26, 27, 28, 29)}, "trn": {y: "TRN" for y in (26, 27, 28, 29)}, "oth": {y: "ASIC" for y in (26, 27, 28, 29)}}
YRS = (26, 27, 28, 29); YC = {26: "C", 27: "D", 28: "E", 29: "F"}

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E）。每列具名範圍＝鍵。金額：每顆加速器美元；市場：十億美元"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源／說明"])
rows = []
sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
sec("A. 加速器顆數（2026–2028 取自 markets/memory；2029 為成長率）")
for k, nm, u, gb, tag, src in MEM.UNITS:
    for i, y in enumerate((26, 27, 28)):
        inp(f"u_{k}_{y}", f"{nm} 顆數 20{y}", "百萬顆", None, u[i], None, tag, ("markets/memory/market.py：" + src) if i == 0 else "同上")
    inp(f"g29_{k}", f"{nm} 2029 顆數成長", "%", None, {"bw": -0.5, "hp": 0.0}.get(k, 0.15), None, "Assumed", "2029 無顆數預測")
sec("B. 每顆加速器網通內容（美元）")
for p, (nm, vals, tag, src) in PLAT.items():
    for (ck, cn), v in zip(COMP, vals):
        inp(f"c_{p}_{ck}", f"{nm}：{cn}", "$/顆", round(v * 0.8), v, round(v * 1.2), tag, src if ck == "su_sw" else "同上")
sec("C. 每顆加速器 IT 資本支出（千美元；算網通 ÷ IT capex）")
for p, (v, tag, src) in ITU.items():
    inp(f"it_{p}", f"{PLAT[p][0]} IT capex", "$K/顆", None, v, None, tag, src)
sec("D. 技術轉換")
for y, v in zip(YRS, (0.01, 0.03, 0.05, 0.08)):
    inp(f"cpo_{y}", f"CPO 占 scale-out 光學金額 20{y}", "%", None, v, None, "Derived／Assumed", "LightCounting：2030 年 CPO 引擎約 $10B vs AI 光學約 $100B（約 10%）；Spectrum-X Photonics 2026-08 量產、放量 2027–28（TrendForce）" if y == 26 else "同上")
for y, v in zip(YRS, (0.0, 0.0, 0.0, 0.2)):
    inp(f"suopt_{y}", f"scale-up 互連中光學占比 20{y}", "%", None, v, None, "Interested-party／Assumed", "NVLink 到 Rubin Ultra Kyber 全銅（The Register）；Feynman「銅與 CPO 都做」（2028 後）" if y == 26 else "同上")
inp("su_cu", "scale-up 互連中銅纜（背板、線纜卡匣）占比", "%", 0.4, 0.6, 0.8, "Assumed", "其餘為交換匣電路板、電源、機構")
sec("E. 廠商歸屬（對帳用，Assumed）")
inp("nv_su_ic", "NVIDIA 平台 scale-up 互連計入 NVIDIA 營收比例", "%", 0.3, 0.5, 0.7, "Assumed", "交換匣含線纜卡匣；背板由 ODM／Amphenol 供應")
inp("nv_so_sw", "NVIDIA 平台 scale-out 交換器中 NVIDIA 份額", "%", 0.35, 0.45, 0.55, "Interested-party／Assumed", "Dell'Oro：Celestica（白牌）第一、NVIDIA 第二、Arista 第三")
inp("nv_opt", "NVIDIA 平台光模組經 NVIDIA（LinkX）銷售比例", "%", 0.1, 0.2, 0.3, "Assumed", "")
inp("avgo_su", "Broadcom 占非 NVIDIA scale-up 交換晶片", "%", 0.3, 0.5, 0.7, "Assumed", "Tomahawk Ultra 2026 Q3 開始部署（Broadcom）")
inp("avgo_so", "Broadcom 晶片占非 NVIDIA scale-out 交換器金額", "%", 0.2, 0.28, 0.35, "Assumed", "交換晶片約占整機 35% × 商用交換晶片市占約 80%")
inp("avgo_nic", "Broadcom 占非 NVIDIA 平台網卡", "%", 0.1, 0.3, 0.5, "Assumed", "")
inp("avgo_opt", "Broadcom 占光學金額（DSP、雷射）", "%", 0.05, 0.10, 0.15, "Assumed", "")
inp("rep_nv26", "NVIDIA 網通營收 2026（實際＋推估）", "$B", 68, 74, 80, "Verified／Derived", "FY27 Q1 14.8（CFO 書面說明）、Q2 約 17.5（法說「季增 18%」）、Q3–Q4 季增 12% Assumed")
inp("rep_avgo26", "Broadcom AI 網通營收 2026", "$B", 16, 17.6, 20, "Derived", "Q2 FY26 AI 10.8 × 約 40%；Q3 16.7 × ≤27%；每季約 4.3–4.5")
inp("rep_opt26", "AI 叢集光學市場 2026（LightCounting）", "$B", None, 26, None, "Interested-party", "LightCounting 2026-01：2025 16.5、2026 26（含 scale-up）")
inp("rep_bern", "網通 ÷ AI 資料中心總 capex（Bernstein）", "%", None, 0.13, None, "Interested-party", "Bernstein 2025-11；IT 約占 75% → ÷ IT 約 17%")
R = {}
r = 4
for s in rows:
    if s[0] == "sec":
        IN.cell(r, 1, s[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, src = s
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None: IN.cell(r, j, v).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, src)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGH", [14, 46, 9, 8, 10, 8, 20, 120]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# ---------------- U：顆數 ----------------
U = wb.create_sheet("Units")
U["A1"] = "Units — 加速器顆數（百萬顆）與對應平台"
hdr(U, 3, ["類型", "2026", "2027", "2028", "2029", "平台（2026／2027+）"])
UR = {}
for i, (k, nm, *_ ) in enumerate(MEM.UNITS):
    rr = 4 + i; UR[k] = rr
    U.cell(rr, 1, nm)
    for y in (26, 27, 28): U.cell(rr, YRS.index(y) + 2, f"=u_{k}_{y}").number_format = "0.00"
    U.cell(rr, 5, f"=D{rr}*(1+g29_{k})").number_format = "0.00"
    U.cell(rr, 6, f"{TYPE2PLAT[k][26]}／{TYPE2PLAT[k][27]}")
UT = 4 + len(MEM.UNITS)
U.cell(UT, 1, "合計").font = BOLD
for j in range(4):
    c = "BCDE"[j]; U.cell(UT, 2 + j, f"=SUM({c}4:{c}{UT-1})").number_format = "0.0"
U.column_dimensions["A"].width = 36; U.column_dimensions["F"].width = 18
UC = {26: "B", 27: "C", 28: "D", 29: "E"}
def units(k, y): return f"Units!{UC[y]}{UR[k]}"

# ---------------- Mkt：市場 ----------------
M = wb.create_sheet("Mkt")
M["A1"] = "Mkt — AI 網通市場（十億美元）＝ Σ 顆數（百萬）× 每顆內容（美元）÷ 1000"
hdr(M, 3, ["項目", "單位", "2026", "2027", "2028", "2029", "說明"])
MR = {}; mr = [4]
def mput(label, unit, fs, key=None, fmt="0.0", note="", bold=False):
    rr = mr[0]
    M.cell(rr, 1, label).font = Font(bold=bold); M.cell(rr, 2, unit)
    for j, f in enumerate(fs):
        if f in (None, ""): continue
        M.cell(rr, 3 + j, f).number_format = fmt
    M.cell(rr, 7, note)
    if key: MR[key] = rr
    mr[0] += 1
def comp_sum(ck, y, filt=None):
    terms = []
    for k, *_ in MEM.UNITS:
        p = TYPE2PLAT[k][y]
        if filt is not None and not filt(p): continue
        terms.append(f"{units(k, y)}*c_{p}_{ck}")
    return "(" + "+".join(terms) + ")/1000" if terms else "0"
NVP = lambda p: p in ("GB", "VR", "RU", "HP")
M.cell(mr[0], 1, "依元件").font = BOLD; mr[0] += 1
for ck, cn in COMP:
    mput(cn, "$B", [f"={comp_sum(ck, y)}" for y in YRS], key=ck)
mput("AI 網通市場合計", "$B", [f"=SUM({YC[y]}{MR['su_sw']}:{YC[y]}{MR['nic']})" for y in YRS], key="tot", bold=True)
mput("  年增", "%", [None] + [f"={YC[y]}{MR['tot']}/{YC[y-1]}{MR['tot']}-1" for y in (27, 28, 29)], fmt="0%")
mput("scale-up（交換＋互連）÷ 合計", "%", [f"=({YC[y]}{MR['su_sw']}+{YC[y]}{MR['su_ic']})/{YC[y]}{MR['tot']}" for y in YRS], key="sush", fmt="0%", note="先驗 N2：Andy 80%")
mput("scale-out（交換、光、銅、網卡）÷ 合計", "%", [f"=1-{YC[y]}{MR['sush']}" for y in YRS], fmt="0%")
mr[0] += 1
M.cell(mr[0], 1, "依產品類別（廠商視角）").font = BOLD; mr[0] += 1
mput("交換晶片與交換器（scale-up＋scale-out）", "$B", [f"={YC[y]}{MR['su_sw']}+{YC[y]}{MR['so_sw']}" for y in YRS], key="sw")
mput("光學（可插拔＋CPO＋scale-up 光學）", "$B", [f"={YC[y]}{MR['so_opt']}+{YC[y]}{MR['su_ic']}*suopt_{y}" for y in YRS], key="opt")
mput("  其中 CPO", "$B", [f"={YC[y]}{MR['so_opt']}*cpo_{y}" for y in YRS], key="cpo")
mput("  其中可插拔", "$B", [f"={YC[y]}{MR['opt']}-{YC[y]}{MR['cpo']}" for y in YRS], key="plug")
mput("銅（scale-up 背板與線纜＋scale-out DAC／AEC）", "$B", [f"={YC[y]}{MR['so_cu']}+{YC[y]}{MR['su_ic']}*(1-suopt_{y})*su_cu" for y in YRS], key="cu")
mput("交換匣其他（電路板、電源、機構）", "$B", [f"={YC[y]}{MR['su_ic']}*(1-suopt_{y})*(1-su_cu)" for y in YRS], key="trayo")
mput("網卡／DPU", "$B", [f"={YC[y]}{MR['nic']}" for y in YRS], key="nic2")
mput("核對：類別合計 − 合計（應為 0）", "$B", [f"={YC[y]}{MR['sw']}+{YC[y]}{MR['opt']}+{YC[y]}{MR['cu']}+{YC[y]}{MR['trayo']}+{YC[y]}{MR['nic2']}-{YC[y]}{MR['tot']}" for y in YRS], fmt="0.000")
mr[0] += 1
M.cell(mr[0], 1, "網通 ÷ AI IT 資本支出").font = BOLD; mr[0] += 1
def it_sum(y):
    return "(" + "+".join(f"{units(k, y)}*it_{TYPE2PLAT[k][y]}" for k, *_ in MEM.UNITS) + ")"
mput("AI IT 資本支出（加速器 × 每顆 IT capex）", "$B", [f"={it_sum(y)}" for y in YRS], key="it", note="百萬顆 × 千美元 ＝ 十億美元")
mput("網通 ÷ AI IT capex", "%", [f"={YC[y]}{MR['tot']}/{YC[y]}{MR['it']}" for y in YRS], key="nit", fmt="0.0%", note="先驗 N1：Andy 2026 15%、2027 20%；Bernstein 換算約 17%")
mput("  NVIDIA 平台 ÷ NVIDIA 平台 IT", "%", ["=(" + "+".join(f"{comp_sum(ck, y, NVP)}" for ck, _ in COMP) + ")/(" + "+".join(f"{units(k, y)}*it_{TYPE2PLAT[k][y]}" for k, *_ in MEM.UNITS if NVP(TYPE2PLAT[k][y])) + ")" for y in YRS], fmt="0.0%")
mr[0] += 1
M.cell(mr[0], 1, "廠商歸屬與對帳（2026）").font = BOLD; mr[0] += 1
def nv_rev(y):
    return (f"={comp_sum('su_sw', y, NVP)}+nv_su_ic*{comp_sum('su_ic', y, NVP)}+nv_so_sw*{comp_sum('so_sw', y, NVP)}"
            f"+{comp_sum('nic', y, NVP)}+nv_opt*{comp_sum('so_opt', y, NVP)}")
NON = lambda p: not NVP(p)
def avgo_rev(y):
    return (f"=avgo_su*{comp_sum('su_sw', y, NON)}+avgo_so*((1-nv_so_sw)*{comp_sum('so_sw', y, NVP)}+{comp_sum('so_sw', y, NON)})"
            f"+avgo_nic*{comp_sum('nic', y, NON)}+avgo_opt*{YC[y]}{MR['opt']}")
mput("NVIDIA 網通（本模型歸屬）", "$B", [nv_rev(y) for y in YRS], key="nv")
mput("  對照：NVIDIA 實際（2026）", "$B", ["=rep_nv26"], key="nvr")
mput("  差距", "%", [f"=C{MR['nv']}/C{MR['nvr']}-1"], fmt="0%")
mput("Broadcom 網通（本模型歸屬）", "$B", [avgo_rev(y) for y in YRS], key="av")
mput("  對照：Broadcom 實際（2026）", "$B", ["=rep_avgo26"], key="avr")
mput("  差距", "%", [f"=C{MR['av']}/C{MR['avr']}-1"], fmt="0%")
mput("NVIDIA＋Broadcom ÷ 市場", "%", [f"=({YC[y]}{MR['nv']}+{YC[y]}{MR['av']})/{YC[y]}{MR['tot']}" for y in YRS], key="nvav", fmt="0%", note="先驗 N3：Andy ≥50%")
mput("光學：本模型 vs LightCounting（2026）", "$B", [f"=C{MR['opt']}", "=rep_opt26"], note="C＝本模型、D＝LightCounting AI 叢集光學（含 scale-up）")
mput("  差距", "%", [f"=C{mr[0]-1}/D{mr[0]-1}-1"], fmt="0%")
M.column_dimensions["A"].width = 44
for c in "CDEF": M.column_dimensions[c].width = 10
M.column_dimensions["G"].width = 60

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入；輸出：2027 AI 網通市場（$B）、網通 ÷ IT capex"
hdr(SE, 3, ["輸入", "2027 市場", "相對基準", "說明"])
base27 = f"Mkt!D{MR['tot']}"
items = [("基準", f"={base27}", ""),
         ("Rubin 顆數 −30%", f"={base27}-0.3*u_rb_27*(" + "+".join(f"c_VR_{ck}" for ck, _ in COMP) + ")/1000", "Rubin 是 2027 最大項"),
         ("Rubin 顆數 +30%", f"={base27}+0.3*u_rb_27*(" + "+".join(f"c_VR_{ck}" for ck, _ in COMP) + ")/1000", ""),
         ("VR200 每顆內容 −20%", f"={base27}-0.2*u_rb_27*(" + "+".join(f"c_VR_{ck}" for ck, _ in COMP) + ")/1000", ""),
         ("VR200 每顆內容 +20%", f"={base27}+0.2*u_rb_27*(" + "+".join(f"c_VR_{ck}" for ck, _ in COMP) + ")/1000", ""),
         ("ASIC（TPU、Trainium、其他）每顆內容 ×2", f"={base27}+(" + "+".join(f"{units(k, 27)}*(" + "+".join(f"c_{TYPE2PLAT[k][27]}_{ck}" for ck, _ in COMP) + ")" for k in ("tpu", "trn", "oth")) + ")/1000", "ASIC 網通內容為 Assumed，不確定性最大")]
for i, (nm, f, note) in enumerate(items):
    rr = 4 + i
    SE.cell(rr, 1, nm); SE.cell(rr, 2, f).number_format = "0.0"
    SE.cell(rr, 3, f"=B{rr}/B$4-1").number_format = "0%"; SE.cell(rr, 4, note)
SE.column_dimensions["A"].width = 40; SE.column_dimensions["D"].width = 40

# ---------------- Sources、README ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源"; SO["A1"].font = BOLD
srcs = [("markets/chip_bridge v0.1（每 GW 網通拆分）", "本 repo"), ("markets/memory/market.py（加速器顆數）", "本 repo"),
 ("NVIDIA Q1 FY27 CFO 書面說明", "https://s201.q4cdn.com/141608511/files/doc_financials/2027/Q127/Q1FY27-CFO-Commentary.pdf"),
 ("NVIDIA Q2 FY27 法說（Webull）", "https://www.webull.com/news/15497781742806016"),
 ("Broadcom Q3 FY26（MarketBeat）", "https://www.marketbeat.com/instant-alerts/transcript-broadcom-q3-earnings-call-highlights-2026-09-02/"),
 ("Dell'Oro 2026-07（2026–30 約 1 兆）", "https://www.delloro.com/news/ai-back-end-switch-sales-to-approach-1-trillion-over-the-next-five-years/"),
 ("Dell'Oro 2Q26", "https://www.delloro.com/news/ai-back-end-networks-switch-sales-surpass-front-end-networks-for-the-first-time-in-2q2026/"),
 ("LightCounting 2026-01 AI 叢集光學", "https://lightcounting.com/newsletter/en/january-2026-optics-for-ai-clusters-366"),
 ("LightCounting 2025-12 CPO", "https://lightcounting.com/newsletter/en/december-2025-aocs-dacs-linear-drive-pluggable-and-co-packaged-optics-320"),
 ("Cignal AI CPO 1Q26", "https://cignal.ai/2026/04/cpo-and-elsfp-1q26-update/"),
 ("TrendForce 2026-07-27（CPO 放量 2027–28）", "https://trendforce.com/presscenter/news/20260727-13151.html"),
 ("The Register 2026-04-05（NVLink 銅／光）", "https://www.theregister.com/2026/04/05/nvidia_optical_scale_up/"),
 ("Bernstein（Investing.com）", "https://www.investing.com/news/stock-market-news/how-much-does-a-gw-of-data-center-capacity-actually-cost-4314046")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 50; SO.column_dimensions["B"].width = 120
RD = wb.create_sheet("README", 0)
lines = [("AI 網通市場模型 v0.1（AI 半導體 Project 步驟 3）　2026-10-08", True),
 ("模型命題：AI 網通的錢由「加速器顆數 × 每顆網通內容」決定；每顆內容隨 scale-up 域擴大（NVL72 → NVL576）與每 GPU 網卡頻寬倍增而上升，而 scale-up 的頻寬大多由便宜的銅承載，所以金額主體仍在 scale-out（交換器、光模組、網卡）。", False), ("", False),
 ("驅動因子 → 推導量", True),
 ("Units：加速器顆數（markets/memory，與 HBM 市場同一份）→ 各年各平台顆數", False),
 ("In B：每顆加速器網通內容（六個元件）；NVIDIA 平台以橋接表 v0.1 為錨，ASIC 為 Assumed", False),
 ("Mkt：Σ 顆數 × 內容 → 依元件、依產品類別的市場；× CPO 占比、scale-up 光學占比 → 光／銅；÷ AI IT capex → 網通占比；× 廠商歸屬 → NVIDIA、Broadcom 推估營收，對照實際", False),
 ("Sens：Rubin 顆數、VR200 內容、ASIC 內容對 2027 市場", False), ("", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。", False),
 ("已知限制：(1) AI 網通總市場與 scale-up／scale-out 美元拆分沒有公開資料（Dell'Oro、650 Group 付費），本模型為由下而上推導；(2) ASIC 網通內容為 Assumed；(3) 廠商歸屬為 Assumed，只用來與 NVIDIA、Broadcom 實際營收對帳；(4) 只含 AI 後端網路，不含前端。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170
wb.save(sys.argv[1]); print("saved", MR)
