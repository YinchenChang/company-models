# 日月光投控（3711.TW；ADR ASX）個股模型 v0.1（AI 半導體 Project：封測次產業）
# 用法：python3 ase/build_ase.py raw.xlsx ；再以 LibreOffice headless 重算
# 樣板：amkr/build_amkr.py（工作表結構、Val 頁列標籤、目標價 DCF＋Comparable、RMX、apply_in_layout）。
# 會計年度＝日曆年，FY(t)＝CY(t)。金額：十億新台幣（NT$B）；LEAP 與資本支出另列美元（US$B）；EPS：新台幣（基本）。
# 口徑：歸屬母公司淨利（IFRS）。行業變數（加速器顆數、情境係數、CoWoS 每片封裝數、CPO 滲透率、匯率）唯讀取自 Siliconomics 快照。
import sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

VER, DATE = "v0.1", "2026-10-31"
BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
SILV = f"Siliconomics 快照 {SIL.VERSION}（{SIL.SNAP['source']['file']}）"

# ---------------- Siliconomics：加速器顆數、平台別、CoWoS 每片封裝數 → 封裝面積指數 ----------------
PLATS = ("bw", "rb", "ru", "hp", "amd", "tpu", "trn", "oth")
UP = {p: SIL.NAMES[f"SIF_UnitsPlat_{p}"] for p in PLATS}          # [None, 2026, 2027, 2028, 2029]
UT = SIL.NAMES["SIF_UnitsPlat_Total_A"]
UNITS = {26: UT[1], 27: UT[2], 28: UT[3], 29: UT[4]}
PPMAP = {"bw": "GB", "rb": "VR", "ru": "RU", "hp": "HP", "amd": ("AMD8", "AMDR"), "tpu": ("TPU7", "TPU8"), "trn": "TRN", "oth": "ASIC"}
def pp(p, y):
    m = PPMAP[p]
    if isinstance(m, tuple): m = m[0] if y == 26 else m[1]
    return SIL.base(f"cw_pp_{m}")
WPU = {}
for j, y in enumerate((26, 27, 28, 29), 1):
    WPU[y] = sum(UP[p][j] / pp(p, y) for p in PLATS) / sum(UP[p][j] for p in PLATS)
AREA = {y: WPU[y] / WPU[26] for y in WPU}

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B），LEAP 與資本支出為十億美元（US$B）；顆數：百萬顆；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = ("標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間（放 Sens），情境專屬列不給區間。SemiAnalysis 未使用。"
            "公司原文＝SEC 6-K（季度新聞稿與簡報）；法說內容經 Motley Fool 逐字稿；公司指引與說法為 Interested-party；券商與媒體為 Interested-party")
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
K3Q = "https://www.sec.gov/Archives/edgar/data/0001122411/000095010326016323/dp255782_6k.htm"
K2Q = "https://www.sec.gov/Archives/edgar/data/0001122411/000095010326011353/dp250875_6k.htm"
K4Q = "https://www.sec.gov/Archives/edgar/data/1122411/000095010326001662/dp241167_6k.htm"
FOOL3Q = "https://www.fool.com/earnings/call-transcripts/2026/10/29/ase-asx-q3-2026-earnings-call-transcript/"
BN2Q = "https://www.bnext.com.tw/article/91677/ase_leap_revenue_2027_capex"
STORM2Q = "https://www.storm.mg/article/11152966"
BOFA = "https://finance.technews.tw/2026/10/14/ase-technology-holding-co-ltd-reported-better-than-expected-results-for-the-third-quarter-of-2026/"
FUBON = "https://fubon-ebrokerdj.fbs.com.tw/Z/ZC/ZCX/ZCX_3711.djhtm"
MT = "https://www.macrotrends.net/stocks/charts/ASX/ase-technology-holding/pe-ratio"
MTROE = "https://www.macrotrends.net/stocks/charts/ASX/ase-technology-holding/roe"
SA = "https://stockanalysis.com/stocks/asx/"
PR3 = "公司 3Q26 新聞稿與簡報（SEC 6-K，2026-10-29）"

sec(f"A. AI 加速器顆數、情境係數與封裝面積（讀 {SILV}；唯讀，不自行重估）")
for y in (26, 27, 28, 29):
    inp(f"units_{y}", f"AI 加速器顆數 20{y}（情境 A 基準）", "百萬顆", None, UNITS[y], None, "Derived",
        "Siliconomics SIF_UnitsPlat_Total_A（八類加速器合計）" if y == 26 else "同上", "NVIDIA、AMD、Google TPU、AWS Trainium、其他 ASIC 合計；與台積電、Amkor、京元電模型同一份" if y == 26 else "同上")
for s in SC:
    for y in (28, 29):
        inp(f"um{y}_{s}", SIL.V[f"um{y}_{s}"]["label"], "x", None, SIL.base(f"um{y}_{s}"), None, "Assumed", f"Siliconomics {f'um{y}_{s}'}（同字母行業情境）",
            "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、Amkor、NVIDIA 模型相同）" if (s == "A" and y == 28) else "同上")
for y in (26, 27, 28, 29):
    inp(f"ar_{y}", f"封裝面積指數 20{y}（加速器用 CoWoS 片數 ÷ 顆數，2026＝1）", "x", None, round(AREA[y], 6), None, "Derived",
        "Siliconomics 平台別顆數 SIF_UnitsPlat_* ÷ 每片封裝數 cw_pp_*（Blackwell→GB、Rubin→VR、Rubin Ultra→RU、AMD 2026 MI3xx→AMD8／之後→AMDR、TPU 2026→TPU7／之後→TPU8）" if y == 26 else "同上",
        "平台對應為本模型假設；每片封裝數下降＝單顆封裝面積變大（Rubin、Rubin Ultra）" if y == 26 else "同上")
for s in SC:
    for y in (26, 27, 28, 29):
        k = f"cpo_{s}_{y}"
        inp(k, SIL.V[k]["label"], "%", None, SIL.base(k), None, "Assumed", f"Siliconomics {k}（技術替代屬行業情境，同字母）", "日月光 CPO 營收依此滲透率相對 2027 情境 A 縮放" if (s == "A" and y == 26) else "同上")
inp("fx", "匯率：新台幣 ÷ 美元（2027 起換算 LEAP 與資本支出）", "NT$/US$", *SIL.vals("fx_twd"), "Assumed", f"Siliconomics fx_twd（{SIL.VERSION}）", "公司 4Q26 指引假設 31.9（3Q26 新聞稿）")

sec("B. 封測（ATM）：LEAP 先進封裝＝顆數 × 委外／份額相對值 × 每顆內容（隨封裝面積）＋全流程（CoWoS full process）＋CPO；主流封測另依成長率")
inp("leap25", "LEAP 營收 FY2025", "US$B", None, 1.6, None, "Interested-party", K4Q + "：LEAP 2025 年 US$1.6B（占 ATM 13%）、2024 年 US$0.6B（公司 4Q25 簡報）", "2026 年初目標 US$3.2B（同上）")
inp("fx25", "FY2025 平均匯率（LEAP 換算新台幣用）", "NT$/US$", None, 31.0, None, "Assumed", "公司：2025 年 ATM 美元營收年增 23%、新台幣年增 20%（4Q25 簡報）；4Q25 指引假設 30.9", "只影響 FY2025 LEAP／主流拆分")
inp("leap26", "LEAP 營收 FY2026（公司：逾 US$4B）", "US$B", None, 4.0, None, "Interested-party", FOOL3Q + "：LEAP 2026 由 US$3.7B 上調到 over US$4 billion（3Q26 法說）", BN2Q + "：2Q26 法說時為逾 US$3.5B")
inp("lsh_1h", "LEAP ÷ ATM 營收 1H26", "%", None, 0.18, None, "Interested-party", BOFA + "：BofA 估 1H26 18%、3Q26 22%、全年約 30%（券商）", "公司未揭露季度 LEAP；只影響 Q 頁季度拆分")
inp("lsh_3q", "LEAP ÷ ATM 營收 3Q26", "%", None, 0.22, None, "Interested-party", "同上（BofA 2026-10-14）", "4Q26＝全年 US$4.0B 減前三季（Derived）")
inp("fp26", "全流程（CoWoS full process）營收 FY2026（含在 LEAP）", "US$B", None, 0.3, None, "Interested-party", FOOL3Q + "：full process 2026 約 US$300M、on track（3Q26 法說）", BN2Q + "：2Q26 法說同（約 3 億美元）")
for s, v, t in (("A", 1.2, "公司稱逾 US$1B 為保守數字；取高於公司說法"), ("B", 1.0, "公司說法下緣（逾 US$1B）"), ("C", 0.7, "全流程導入延後、台積電自製比例維持")):
    inp(f"fp27_{s}", f"情境 {s}：全流程營收 FY2027", "US$B", None, v, None, "Interested-party" if s == "B" else "Assumed",
        FOOL3Q + "：full process 2027 over US$1 billion（保守；3Q26 法說）" if s == "A" else "同上", t)
inp("fp_g", "全流程 FY2028–29：每年隨顆數成長外再加的滲透成長", "%", 0.1, 0.3, 0.5, "Assumed", "台積電 CoWoS 2027 需求接近產能（Siliconomics cw_capend27 200 千片／月）；全流程擴大由 OSAT 補充", "量化資料：未找到")
inp("cpo27", "CPO 營收 FY2027（情境 A 滲透率下；含在 LEAP）", "US$B", 0.3, 0.5, 0.8, "Interested-party", FOOL3Q + "：CPO 2027 over US$0.5 billion，含在 LEAP；4Q26 量產（3Q26 法說）", BN2Q + "：2Q26 時稱 2026 年底小量（媒體）")
for s, a, b, c in (("A", 1.10, 1.15, 1.20, ), ("B", 1.05, 1.08, 1.10), ("C", 1.00, 0.95, 0.92)):
    for y, v in ((27, a), (28, b), (29, c)):
        inp(f"rel{y}_{s}", f"情境 {s}：日月光委外與份額相對值 20{y}（2026＝1）", "x", None, v, None, "Assumed",
            "台積電 CoWoS 吃緊、委外補充（Siliconomics cw_osat 10%）；日月光直接專案擴大（AMD Venice 等；BofA 2026-10-14）" if (s == "A" and y == 27) else "同上",
            "A＝委外擴大、份額上升；B＝小幅上升（公司 LEAP 翻倍的中性路徑）；C＝台積電自製與 Amkor 分食" if (s == "A" and y == 27) else "同上")
inp("el", "每顆 LEAP 內容對封裝面積的彈性", "x", 0.3, 0.5, 1.0, "Assumed", "封裝面積大 → oS、測試時間與良率成本上升；後段金額不與面積等比例（與 Amkor 模型同值）", "第二來源：未找到")
inp("g_ms26", "主流封測年增 FY2026（參考；錨點已由 Q 頁實績決定）", "%", None, "=Q!G7/Q!B7-1", None, "Derived", "Q 頁：FY2026 錨點 ÷ FY2025（主流＝ATM − LEAP）", FOOL3Q + "：公司稱 2026 年 general market 成長逾 20%")
inp("g_ms27", "主流封測年增 FY2027", "%", 0.05, 0.15, 0.22, "Interested-party", FOOL3Q + "：公司預期 2027 年 general market 成長與 2026 相近（逾 20%；3Q26 法說）", "本模型基準取 15%（漲價與成熟製程滿載不必然延續）；行業變數，待併入 Siliconomics（非 AI 封測需求）")
inp("g_msL", "主流封測年增 FY2028–29", "%", 0.02, 0.07, 0.12, "Assumed", "非 AI 半導體封測長期成長；公司 3Q26：運算 +39%、車用工業 +14%、消費 +14%、通訊 +8%（年增）", "行業變數，待併入 Siliconomics（非 AI 封測需求）")

sec("C. EMS（環旭電子；簡單成長率）與沖銷")
inp("g_ems27", "EMS 年增 FY2027", "%", 0.0, 0.10, 0.15, "Assumed", FOOL3Q + "：EMS 2027 成長，由 HPC 帶動（約占 EMS mid-teens、快速放量）；通訊、消費較弱（3Q26 法說）", "量化指引：未找到")
inp("g_emsL", "EMS 年增 FY2028–29", "%", 0.0, 0.05, 0.08, "Assumed", "FY2024→FY2025 EMS −5%（公司）；2026 年 2H 含 HPC 新案", "—")
inp("ems_gm", "EMS 毛利率（FY2027 起）", "%", 0.075, 0.085, 0.095, "Derived", "FY2025 9.1%、1Q26 9.5%、2Q26 8.9%、3Q26 8.3%（公司分部表）", "4Q26 指引 EMS 營業利益率約 2%（3Q26 3.3%）")

sec("D. 資本支出與折舊（資本支出約占營收四成：核心問題之二）")
inp("capex25", "資本支出 FY2025", "US$B", None, 5.5, None, "Interested-party", K4Q + "：機器設備 US$3.4B＋廠房設施與自動化 US$2.1B（公司 4Q25 簡報）", "—")
inp("capex26", "資本支出 FY2026（公司：US$11.5B）", "US$B", None, 11.5, None, "Interested-party", FOOL3Q + "：2026 資本支出由 US$10.5B 上調到 US$11.5B；前三季 US$8B（設備 4.8、廠房 3.2）", BN2Q + "：2Q26 時為 US$10.5B（廠房 4.0、設備 6.5）")
inp("capex27", "資本支出 FY2027", "US$B", 11.5, 13.0, 15.0, "Interested-party", FOOL3Q + "：2027 資本支出 even larger than this year；2027 現金流為負（3Q26 法說）", BOFA + "：BofA 估 2027–28 每年逾 US$13B（券商）")
inp("capi", "資本支出 ÷ 合併營收（FY2028–29）", "%", 0.22, 0.30, 0.38, "Assumed", "FY2025 US$5.5B × 31.0 ÷ 645.4＝26%；FY2026 約 42%（Q 頁）", "LEAP 成長放緩後回落（Assumed）")
inp("life", "新增資本支出的平均折舊年限（廠房與設備混合）", "年", 6, 8, 12, "Assumed", "2026 年資本支出約 35% 為廠房設施（公司）；設備 5–7 年、廠房 20 年以上", "混合年限：未找到")
inp("q4_da", "4Q26 折舊攤提", "NT$B", None, 27.0, None, "Assumed", "1Q26 20.7、2Q26 24.6、3Q26 24.8（EBITDA − 營業利益；公司簡報推得）；3Q26 資本支出 US$3.9B", "—")

sec("E. 毛利率、營業費用與淨利（FY2027 起）")
inp("ft", "封測現金毛利增量 ÷ 封測營收增量（現金流通率）", "%", 0.40, 0.50, 0.60, "Assumed", "FY2025→FY2026 錨點實測見 Q 頁（漲價、LEAP 組合、稼動率滿載）", FOOL3Q + "：公司稱 30% 毛利率由天花板變成基準、2027 逐季改善；BofA 估 AI／HPC 專案毛利率 35–45%")
inp("phi", "營業費用成長對營收成長的彈性（FY2027 起）", "x", 0.3, 0.6, 1.0, "Assumed", "FY2026 實測見 Q 頁", "公司：增加研發、人才與智慧工廠投入（2Q26 簡報）")
inp("nonop_f", "營業外淨額（FY2027 起每年）", "NT$B", -8.0, -5.0, -2.0, "Assumed", "1Q26 +0.7、2Q26 +4.6（含投資評價利益 4.2、匯兌避險 1.5，淨利息 −1.9；媒體）、3Q26 −0.4（公司）", "淨負債 ÷ 權益 3Q26 66%（2Q26 47%），2027 現金流為負（公司）")
inp("tax", "有效稅率（FY2027 起）", "%", 0.15, 0.17, 0.20, "Interested-party", FOOL3Q + "：2026 全年有效稅率 16–17%（3Q26 法說）", "FY2025 18.4%（公司損益推得）")
inp("nci", "非控制權益（FY2027 起每年）", "NT$B", None, 3.0, None, "Derived", "FY2025 1.2、1Q–3Q26 0.4／0.5／0.9（公司）", "—")
inp("sh_g", "基本股數年變動", "%", None, 0.005, None, "Assumed", "FY2025 4,339 → 3Q26 4,404 百萬股（淨利 ÷ 基本 EPS 推得）；員工認股與可轉債轉換", "—")

sec("F. FY2026 錨點（1Q–3Q26 實績＋4Q26 依公司指引）")
inp("q4_atm", "4Q26 封測（ATM）季增", "%", None, 0.12, None, "Interested-party", PR3 + "：4Q26 ATM 營收季增 11–13%（新台幣）；取中點", "—")
inp("q4_atmgm", "4Q26 封測毛利率增幅（百分點）", "%", None, 0.015, None, "Interested-party", PR3 + "：4Q26 ATM 毛利率季增 100–200 bps；取中點", "—")
inp("q4_ems", "4Q26 EMS 季增", "%", None, 0.0, None, "Interested-party", PR3 + "：4Q26 EMS 營收約持平；營業利益率約 2%", "—")
inp("q4_emsgm", "4Q26 EMS 毛利率", "%", None, 0.080, None, "Assumed", "3Q26 8.3%；營業利益率指引由 3.3% 降到約 2%", "—")
inp("q4_opx", "4Q26 營業費用季增", "%", None, 0.08, None, "Assumed", "3Q26 營業費用 20.8（毛利 − 營業利益）；以合併營業利益率指引 14.5–15.5% 校準", "—")
inp("q4_nonop", "4Q26 營業外淨額", "NT$B", None, -0.5, None, "Assumed", "3Q26 −0.4；淨負債上升", "—")
inp("q4_tax", "4Q26 有效稅率", "%", None, 0.17, None, "Interested-party", "公司全年 16–17%（3Q26 法說）", "3Q26 14.8%（公司損益推得）")
inp("q4_nci", "4Q26 非控制權益", "NT$B", None, 0.8, None, "Derived", "3Q26 0.9（公司）", "—")

sec("G. 公司指引、券商預估（只作 Recon 對帳，不作校準）")
inp("g4_rev_lo", "4Q26 合併營收季增指引下緣", "%", None, 0.06, None, "Interested-party", PR3, "上緣 8%")
inp("g4_rev_hi", "4Q26 合併營收季增指引上緣", "%", None, 0.08, None, "Interested-party", PR3, "—")
inp("g4_gm_lo", "4Q26 合併毛利率指引下緣", "%", None, 0.23, None, "Interested-party", PR3, "上緣 24%")
inp("g4_gm_hi", "4Q26 合併毛利率指引上緣", "%", None, 0.24, None, "Interested-party", PR3, "—")
inp("g4_om_lo", "4Q26 合併營業利益率指引下緣", "%", None, 0.145, None, "Interested-party", PR3, "上緣 15.5%")
inp("g4_om_hi", "4Q26 合併營業利益率指引上緣", "%", None, 0.155, None, "Interested-party", PR3, "—")
inp("bofa_eps26", "BofA EPS FY2026", "NT$", None, 19.64, None, "Interested-party", BOFA + "：BofA 2026-10-14（3Q26 財報前、月營收後；券商）", "賣方共識：未找到（FactSet、stockanalysis 皆未取得）")
inp("bofa_eps27", "BofA EPS FY2027", "NT$", None, 34.26, None, "Interested-party", "同上", "—")
inp("bofa_eps28", "BofA EPS FY2028", "NT$", None, 49.78, None, "Interested-party", "同上", "—")
inp("bofa_gm28", "BofA 封測毛利率 FY2028（約 30% 中段）", "%", None, 0.30, None, "Interested-party", "同上：2026 29%、2027 約 30%（低段）、2028 約 30%（中段）", "—")

sec("H. 價格與市場倍數")
inp("px", "股價（2026-10-30 收盤，臺灣證券交易所）", "NT$", None, 668.0, None, "Interested-party", FUBON + "：10-30 收盤 668.00（−15.00）", "第二來源：未找到（Yahoo 奇摩股市頁面為 10-27 盤中 652）；ADR ASX 10-26 收盤 US$39.91（stockanalysis；1 ADS＝2 股）")
inp("pe_ttm", "近 12 月 P/E（Yahoo 奇摩股市，10-27 盤中）", "x", None, 47.32, None, "Interested-party", "https://tw.stock.yahoo.com/quote/3711", "stockanalysis ADR 49.19（10-26）")
inp("fpe", "前瞻 P/E（stockanalysis，ADR，10-26）", "x", None, 26.99, None, "Interested-party", SA, "3Q26 財報前")
inp("cons_feps", "共識前瞻 12 月 EPS（BofA：2／12 × FY2027＋10／12 × FY2028；單一券商）", "NT$", None, "=2/12*bofa_eps27+10/12*bofa_eps28", None, "Derived", "BofA 2026-10-14（單一券商；與本模型 12 個月後前瞻 EPS 同口徑）", "賣方共識前瞻 EPS：未找到")
inp("cons_tp", "市場共識目標價（以 BofA 單一券商代替）", "NT$", None, 900.0, None, "Interested-party", BOFA + "：BofA 目標價 800 → 900（20 倍目標 P/E；2026-10-14）", "賣方平均目標價：未找到；經濟日報 2026-10-16：法人目標價調升近一成（無數字）")
sec("I. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%；各市場同一值）", "DCF 與 Sens")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "成熟倍數與 FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均，FY2023–25）", "%", None, "=AVERAGE(0.1172,0.0973,0.1015)", None, "Derived", MTROE + "：2023 11.72%、2024 9.73%、2025（近 12 月至 2025-03）10.15%", "單一來源；2026 年 ROE 隨獲利倍增明顯上升（未納入）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（工作單 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2029 淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（情境 B FY2029 與其 × 0.8 的中點）", "%", None, 0, None, "Assumed", "十年平均淨利率未找到（FY2016–2023 本專案未逐年核對）：工作單規則以 FY2029 × 0.8 代替", "PL 頁情境 B FY2029（公式於建置時連結）")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（＝情境 C FY2029 × 0.8）", "%", None, 0, None, "Assumed", "同上（十年平均未找到）", "PL 頁情境 C FY2029（公式於建置時連結）")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

URLNAME = {K3Q: "公司 3Q26 新聞稿與簡報（SEC 6-K）", K2Q: "公司 2Q26 新聞稿與簡報（SEC 6-K）", K4Q: "公司 4Q25 新聞稿與簡報（SEC 6-K）",
           FOOL3Q: "3Q26 法說逐字稿（Motley Fool）", BN2Q: "2Q26 法說報導（數位時代）", STORM2Q: "2Q26 法說報導（風傳媒）",
           BOFA: "BofA 報告摘要（科技新報 2026-10-14）", FUBON: "富邦證券 e 點通（3711 日線）", MT: "macrotrends", MTROE: "macrotrends ROE", SA: "stockanalysis.com"}
r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    for u_, nm_ in URLNAME.items():
        s1 = s1.replace(u_, nm_); s2 = s2.replace(u_, nm_)
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for k in TP_KEYS:
    for c in range(1, 10): IN.cell(INROW[k], c).fill = KEY
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點（NT$B） ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣、基本；股數為百萬股）。藍字＝輸入（SEC 6-K 公司新聞稿與簡報，原單位百萬元 ÷ 1,000）；黑字＝公式。分部營收取分部表（含部門間交易），差額列「沖銷及其他」"
Q["A2"] = "v0.1（2026-10-31）：FY2026 錨點＝1Q、2Q、3Q26 實績（3Q26：2026-10-29 新聞稿）＋4Q26 依公司指引（ATM 季增 12%、ATM 毛利率 +1.5pt、EMS 持平；合併營收季增約 7%，指引 6–8%）"
hdr(Q, 3, ["項目", "FY2025A", "1Q26（實績）", "2Q26（實績）", "3Q26（實績）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：4Q25 新聞稿（SEC 6-K 2026-02-05）；1Q、2Q26：2Q26 新聞稿（6-K 2026-07-30）；3Q26：6-K 2026-10-29"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2="—", ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
qrow(4, "營收", {"B": 645.388, "C": 173.662, "D": 191.064, "E": 244.161, "F": "=F5+F8+F9"}, s2="9M26 608.887（公司月營收）＝三季合計")
qrow(5, "封測（ATM；分部表）", {"B": 389.228, "C": 112.434, "D": 126.148, "E": 151.485, "F": "=E5*(1+q4_atm)"}, tag="Verified／Interested-party（4Q）", s2="ATM＝封裝、測試、材料；合併表口徑 FY2025 385.3、3Q26 150.6（含沖銷）")
qrow(6, "  LEAP 先進封裝（命題子列）", {"B": "=leap25*fx25", "C": "=C5*lsh_1h", "D": "=D5*lsh_1h", "E": "=E5*lsh_3q", "F": "=leap26*31.9-SUM(C6:E6)"}, tag="Interested-party／Derived",
     s2="FY2025 US$1.6B；1H26、3Q26 依 BofA 占比；4Q26＝全年 US$4.0B × 31.9 − 前三季（Derived）")
qrow(7, "  主流封測（ATM − LEAP）", {c: f"={c}5-{c}6" for c in "BCDEF"}, tag="Derived", s1="—")
qrow(8, "EMS（環旭；分部表）", {"B": 259.079, "C": 61.875, "D": 65.789, "E": 93.469, "F": "=E8*(1+q4_ems)"}, tag="Verified／Interested-party（4Q）", s2="1Q26 簡報另一頁 61.361（口徑不同）；取分部頁")
qrow(9, "沖銷及其他（合併 − 分部合計）", {"B": "=B4-B5-B8", "C": "=C4-C5-C8", "D": "=D4-D5-D8", "E": "=E4-E5-E8", "F": "=E9/(E5+E8)*(F5+F8)"}, tag="Derived", s1="—", s2="4Q26 依 3Q26 比例")
qrow(10, "封測毛利", {"B": 91.380, "C": 29.198, "D": 34.487, "E": 46.412, "F": "=F5*F11"}, tag="Verified")
qrow(11, "封測毛利率", {"B": "=B10/B5", "C": "=C10/C5", "D": "=D10/D5", "E": "=E10/E5", "F": "=E11+q4_atmgm", "G": "=G10/G5"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="公司：FY2025 23.5%、3Q26 30.6%")
qrow(12, "EMS 毛利", {"B": 23.695, "C": 5.862, "D": 5.873, "E": 7.793, "F": "=F8*q4_emsgm"}, tag="Verified")
qrow(13, "沖銷及其他毛利（合併毛利 − 分部毛利）", {"B": "=B14-B10-B12", "C": "=C14-C10-C12", "D": "=D14-D10-D12", "E": "=E14-E10-E12", "F": "=E13/E9*F9"}, tag="Derived", s1="—")
qrow(14, "毛利", {"B": 114.193, "C": 34.818, "D": 40.147, "E": 53.967, "F": "=F10+F12+F13"}, tag="Verified")
qrow(15, "毛利率", {c: f"={c}14/{c}4" for c in "BCDEFG"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="公司：FY2025 17.7%、3Q26 22.1%")
qrow(16, "營業費用（毛利 − 營業利益）", {"B": "=B14-B17", "C": "=C14-C17", "D": "=D14-D17", "E": "=E14-E17", "F": "=E16*(1+q4_opx)"}, tag="Derived", s1="—")
qrow(17, "營業利益", {"B": 50.756, "C": 17.493, "D": 21.130, "E": 33.142, "F": "=F14-F16"}, tag="Verified")
qrow(18, "營業外淨額", {"B": 0.545, "C": 0.668, "D": 4.566, "E": "=E19-E17", "F": "=q4_nonop"}, tag="Verified／Derived", s2="2Q26 含投資評價利益 4.2、匯兌避險 1.5（風傳媒）")
qrow(19, "稅前利益", {"B": 51.301, "C": 18.161, "D": 25.696, "E": 32.710, "F": "=F17+F18"}, tag="Verified")
qrow(20, "所得稅", {"B": 9.461, "C": 3.629, "D": 4.171, "E": 4.831, "F": "=F19*q4_tax"}, tag="Verified")
qrow(21, "非控制權益", {"B": 1.182, "C": 0.400, "D": 0.460, "E": 0.883, "F": "=q4_nci"}, tag="Verified")
qrow(22, "歸屬母公司淨利", {"B": 40.658, "C": 14.132, "D": 21.065, "E": 26.996, "F": "=F19-F20-F21"}, tag="Verified")
qrow(23, "基本股數（百萬股）", {"B": "=B22/B25*1000", "C": "=C22/C25*1000", "D": "=D22/D25*1000", "E": "=E22/E25*1000", "F": "=E23"}, fmt="#,##0.0", ann="=AVERAGE(C23:F23)", tag="Derived", s1="淨利 ÷ 公司揭露基本 EPS")
qrow(24, "EPS（基本）", {c: f"={c}22/{c}23*1000" for c in "BCDEF"}, fmt="0.00", ann="=SUM(C24:F24)", tag="Derived", s1="淨利 ÷ 股數；FY2026＝四季加總")
qrow(25, "  公司揭露 EPS（基本）", {"B": 9.37, "C": 3.23, "D": 4.80, "E": 6.13}, fmt="0.00", ann="", tag="Verified", s2="稀釋 EPS：FY2025 8.89、3Q26 5.84（可轉債）")
qrow(26, "折舊攤提（EBITDA − 營業利益）", {"B": "=126.011-B17", "C": "=38.147-C17", "D": "=45.779-D17", "E": "=57.963-E17", "F": "=q4_da"}, tag="Derived", s2="EBITDA：FY2025 126.0、1Q26 38.1、2Q26 45.8、3Q26 58.0（公司簡報）")
qrow(27, "資本支出（US$B）", {"B": "=capex25", "G": "=capex26"}, fmt="0.00", ann="", tag="Interested-party", s2="公司：2026 前三季 US$8.0B；全年 US$11.5B")
qrow(28, "資本支出 ÷ 營收", {"B": "=B27*fx25/B4", "G": "=G27*31.9/G4"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2026 以 31.9 換算")
qrow(29, "LEAP ÷ 封測營收", {c: f"={c}6/{c}5" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—")
Q.cell(31, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
checks = [("FY2026 營收年增", "=G4/B4-1", "0.0%"), ("FY2026 封測年增（公司 2Q26：逾 35%）", "=G5/B5-1", "0.0%"), ("FY2026 主流封測年增（公司：逾 20%）", "=G7/B7-1", "0.0%"),
          ("封測現金毛利增量 ÷ 封測營收增量（In 頁 ft 的對照；折舊增量全數歸封測）", "=((G10+G26)-(B10+B26))/(G5-B5)", "0.0%"),
          ("營業費用彈性（In 頁 phi 的對照）", "=(G16/B16-1)/(G4/B4-1)", "0.00"),
          ("4Q26 合併營收季增（指引 6–8%）", "=F4/E4-1", "0.0%"), ("4Q26 合併毛利率（指引 23–24%）", "=F15", "0.0%"), ("4Q26 合併營業利益率（指引 14.5–15.5%）", "=F17/F4", "0.0%"),
          ("FY2025 分部毛利 + 沖銷 − 合併毛利（應為 0）", "=B10+B12+B13-B14", "0.000"), ("3Q26 EPS 自算 − 公司 6.13（應約 0）", "=E24-E25", "0.000")]
for i, (lab, f, fm) in enumerate(checks, 32):
    Q.cell(i, 1, lab); Q[f"B{i}"] = f; Q[f"B{i}"].number_format = fm
for nm, rr in (("rev", 4), ("atm", 5), ("leap", 6), ("ms", 7), ("ems", 8), ("elim", 9), ("agp", 10), ("egp", 12), ("lgp", 13), ("gp", 14), ("opx", 16), ("oi", 17), ("nonop", 18), ("tax", 20), ("nci", 21), ("ni", 22), ("sh", 23), ("eps", 24), ("da", 26)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 62
for c in "BCDEFG": Q.column_dimensions[c].width = 14
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 60; Q.column_dimensions["J"].width = 60
Q.freeze_panes = "B4"

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I", 2032: "J"}
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
BU = (2026, 2027, 2028, 2029); F3 = (2027, 2028, 2029)
DEFP = dict(scn="A", mf="1", rel="rel{y}_{s}", el="el", fp27="fp27_{s}", fpg="fp_g", cpo27="cpo27", cpo="cpo_{s}_{y}",
            gms27="g_ms27", gmsL="g_msL", gems27="g_ems27", gemsL="g_emsL", emsgm="ems_gm", fx="fx",
            capex27="capex27", capi="capi", life="life", ft="ft", phi="phi", nonop="nonop_f", tax="tax", fg="fade_g{s}", fm="fade_m{s}")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    p = lambda k, y=None: P[k].format(y=(y % 100 if y else ""), s=s)
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 11, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    um = lambda y: {2026: "", 2027: "", 2028: f"*um28_{s}", 2029: f"*um29_{s}"}[y]
    sec_("A. AI 加速器 → LEAP 先進封裝（D 欄＝FY2026 校準年；美元）")
    row("units", "AI 加速器顆數（含情境係數）", "百萬顆", {FC[y]: f"=units_{y % 100}{um(y)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU}, "0.00")
    row("area", "封裝面積指數（2026＝1）", "x", {FC[y]: f"=ar_{y % 100}" for y in BU}, "0.000")
    row("rel", "委外與份額相對值（2026＝1）", "x", {"D": "=1", **{FC[y]: f"={p('rel', y)}" for y in F3}}, "0.00")
    row("v", "每顆加速器 LEAP 核心內容（不含全流程、CPO）", "US$／顆", {"D": f"=D{r + 1}*1000/D{R['units']}", **{FC[y]: f"=$D{r}*{FC[y]}{R['area']}^{P['el']}" for y in F3}}, "#,##0",
        note="FY2026 校準（Derived）；之後 × 封裝面積指數 ^ 彈性")
    row("core", "LEAP 核心（2.5D oS、先進測試等）", "US$B", {"D": f"=leap26-D{r + 1}-D{r + 2}", **{FC[y]: f"={FC[y]}{R['units']}*{FC[y]}{R['rel']}*{FC[y]}{R['v']}/1000" for y in F3}}, "0.000",
        note="FY2026＝公司 LEAP 逾 US$4B − 全流程 − CPO；之後＝顆數 × 相對值 × 每顆內容")
    row("fp", "全流程（CoWoS full process）", "US$B", {"D": "=fp26", "E": f"={p('fp27')}", **{FC[y]: f"={FC[y - 1]}{r}*{FC[y]}{R['units']}/{FC[y - 1]}{R['units']}*(1+{P['fpg']})" for y in (2028, 2029)}}, "0.000")
    row("cpo", "CPO（共同封裝光學）", "US$B", {FC[y]: f"={P['cpo27']}*{p('cpo', y)}/cpo_A_27" + (f"*{FC[y]}{R['units']}/E{R['units']}" if y > 2027 else "") for y in BU}, "0.000",
        note="2027 情境 A＝公司逾 US$0.5B；依 Siliconomics 同字母 CPO 滲透率縮放（2028 起另乘顆數比）")
    row("leapu", "LEAP 合計", "US$B", {"C": "=leap25", **{FC[y]: f"={FC[y]}{R['core']}+{FC[y]}{R['fp']}+{FC[y]}{R['cpo']}" for y in BU}}, "0.000")
    row("fxr", "匯率", "NT$/US$", {"C": "=fx25", "D": "=a26_leap/D" + str(r - 1), **{FC[y]: f"={P['fx']}" for y in F3}}, "0.00", note="FY2026＝Q 頁 LEAP 新台幣 ÷ 美元（隱含）")
    sec_("B. 營收分部（十億新台幣）")
    row("leap", "  LEAP 先進封裝（命題子列）", "NT$B", {"C": "=a25_leap", "D": "=a26_leap", **{FC[y]: f"={FC[y]}{R['leapu']}*{FC[y]}{R['fxr']}" for y in F3}})
    row("ms", "  主流封測", "NT$B", {"C": "=a25_ms", "D": "=a26_ms", "E": f"=D{r}*(1+{P['gms27']})", **{FC[y]: f"={FC[y - 1]}{r}*(1+{P['gmsL']})" for y in (2028, 2029)}})
    row("atm", "封測（ATM）", "NT$B", {FC[y]: f"={FC[y]}{R['leap']}+{FC[y]}{R['ms']}" for y in FYS})
    row("ems", "EMS（環旭）", "NT$B", {"C": "=a25_ems", "D": "=a26_ems", "E": f"=D{r}*(1+{P['gems27']})", **{FC[y]: f"={FC[y - 1]}{r}*(1+{P['gemsL']})" for y in (2028, 2029)}})
    row("elim", "沖銷及其他", "NT$B", {"C": "=a25_elim", "D": "=a26_elim", **{FC[y]: f"=$D{r}/($D{R['atm']}+$D{R['ems']})*({FC[y]}{R['atm']}+{FC[y]}{R['ems']})" for y in F3}}, note="依 FY2026 比例")
    row("rev", "營收", "NT$B", {FC[y]: f"={FC[y]}{R['atm']}+{FC[y]}{R['ems']}+{FC[y]}{R['elim']}" for y in FYS})
    row("g", "營收年增", "%", {FC[y]: f"={FC[y]}{R['rev']}/{FC[y - 1]}{R['rev']}-1" for y in FYS[1:]}, "0.0%")
    row("lsh", "LEAP ÷ 封測營收", "%", {FC[y]: f"={FC[y]}{R['leap']}/{FC[y]}{R['atm']}" for y in FYS}, "0.0%")
    row("atmg", "封測年增", "%", {FC[y]: f"={FC[y]}{R['atm']}/{FC[y - 1]}{R['atm']}-1" for y in FYS[1:]}, "0.0%")
    sec_("C. 資本支出、折舊、毛利與損益")
    row("capu", "資本支出", "US$B", {"C": "=capex25", "D": "=capex26", "E": f"={P['capex27']}", **{FC[y]: f"={P['capi']}*{FC[y]}{R['rev']}/{FC[y]}{R['fxr']}" for y in (2028, 2029)}}, "0.00",
        note="FY2026 公司 US$11.5B；FY2027 公司「更大」；之後＝營收 × 資本支出率")
    row("cap", "資本支出（新台幣）", "NT$B", {FC[y]: f"={FC[y]}{R['capu']}*" + ("31.9" if y == 2026 else f"{FC[y]}{R['fxr']}") for y in FYS})
    row("capr", "資本支出 ÷ 營收", "%", {FC[y]: f"={FC[y]}{R['cap']}/{FC[y]}{R['rev']}" for y in FYS}, "0.0%")
    row("da", "折舊攤提", "NT$B", {"C": "=a25_da", "D": "=a26_da", **{FC[y]: f"={FC[y - 1]}{r}*(1-1/{P['life']})+{FC[y - 1]}{R['cap']}/{P['life']}" for y in F3}},
        note="前一年 ×（1 − 1／年限）＋前一年資本支出 ÷ 年限")
    row("agp", "封測毛利 ＝ FY2026 ＋ 流通率 × 封測營收增量 − 折舊增量", "NT$B",
        {"C": "=a25_agp", "D": "=a26_agp", **{FC[y]: f"=$D{r}+{P['ft']}*({FC[y]}{R['atm']}-$D{R['atm']})-({FC[y]}{R['da']}-$D{R['da']})" for y in F3}}, note="流通率為現金毛利口徑；折舊增量全數歸封測")
    row("agm", "封測毛利率", "%", {FC[y]: f"={FC[y]}{R['agp']}/{FC[y]}{R['atm']}" for y in FYS}, "0.00%")
    row("egp", "EMS 毛利", "NT$B", {"C": "=a25_egp", "D": "=a26_egp", **{FC[y]: f"={FC[y]}{R['ems']}*{P['emsgm']}" for y in F3}})
    row("lgp", "沖銷及其他毛利", "NT$B", {"C": "=a25_lgp", "D": "=a26_lgp", **{FC[y]: f"=$D{r}/$D{R['elim']}*{FC[y]}{R['elim']}" for y in F3}})
    row("gp", "毛利", "NT$B", {FC[y]: f"={FC[y]}{R['agp']}+{FC[y]}{R['egp']}+{FC[y]}{R['lgp']}" for y in FYS})
    row("gm", "毛利率", "%", {FC[y]: f"={FC[y]}{R['gp']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("opx", "營業費用", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{FC[y]: f"={FC[y - 1]}{r}*(1+{P['phi']}*{FC[y]}{R['g']})" for y in F3}})
    row("oi", "營業利益", "NT$B", {FC[y]: f"={FC[y]}{R['gp']}-{FC[y]}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {FC[y]: f"={FC[y]}{R['oi']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("nonop", "營業外淨額", "NT$B", {"C": "=a25_nonop", "D": "=a26_nonop", **{FC[y]: f"={P['nonop']}" for y in F3}})
    row("pti", "稅前利益", "NT$B", {FC[y]: f"={FC[y]}{R['oi']}+{FC[y]}{R['nonop']}" for y in FYS})
    row("tax", "所得稅", "NT$B", {"C": "=a25_tax", "D": "=a26_tax", **{FC[y]: f"={FC[y]}{R['pti']}*{P['tax']}" for y in F3}})
    row("nci", "非控制權益", "NT$B", {"C": "=a25_nci", "D": "=a26_nci", **{FC[y]: "=nci" for y in F3}})
    row("ni", "歸屬母公司淨利", "NT$B", {FC[y]: f"={FC[y]}{R['pti']}-{FC[y]}{R['tax']}-{FC[y]}{R['nci']}" for y in FYS})
    row("shr", "基本股數", "百萬股", {"C": "=a25_sh", "D": "=a26_sh", **{FC[y]: f"={FC[y - 1]}{r}*(1+sh_g)" for y in F3}}, "#,##0.0")
    row("eps", "EPS（基本）", "NT$", {"C": "=a25_eps", "D": "=a26_eps", **{FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['shr']}*1000" for y in F3}}, "0.00")
    row("epsg", "EPS 年增", "%", {FC[y]: f"={FC[y]}{R['eps']}/{FC[y - 1]}{R['eps']}-1" for y in FYS[1:]}, "0.0%")
    row("nim", "淨利率（歸屬母公司淨利 ÷ 營收）", "%", {FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("fcfe", "股東自由現金流（淨利＋折舊 − 資本支出；參考）", "NT$B", {FC[y]: f"={FC[y]}{R['ni']}+{FC[y]}{R['da']}-{FC[y]}{R['cap']}" for y in FYS}, note="不含營運資金與融資；只作參考，不進 DCF")
    row("roic", "營業利益增量（相對 FY2025）÷ FY2026 起累計資本支出", "%", {FC[y]: f"=({FC[y]}{R['oi']}-$C{R['oi']})/SUM($D{R['cap']}:{FC[y]}{R['cap']})" for y in BU}, "0.0%",
        note="資本支出的稅前邊際報酬（核心問題之二）")
    for y in FADE:
        c, pc, k = FC[y], FC[y - 1], y - 2029
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+({p('fg')}-$G{R['g']})*{k}/3"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({p('fm')}-$G{R['nim']})*{k}/3"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['shr']}"] = f"={pc}{R['shr']}*(1+sh_g)"
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['shr']}*1000"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("g", "0.0%"), ("rev", "#,##0.00"), ("nim", "0.00%"), ("ni", "#,##0.00"), ("shr", "#,##0.0"), ("eps", "0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fm; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 11, "FY2030 起為收斂法：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導")
    R["end"] = r
    return R

def dcf_expr(cells, r="tp_r", g="tp_g"):
    """一個情境的 DCF 每股價值（股東現金流＝EPS ×（1 − g／ROE）；下限 0）。cells：FY2027–FY2032 EPS 儲存格（6 個）。"""
    reinv = f"MIN({g}/tp_roe,0.8)"
    cf = lambda x: f"{x}*(1-{reinv})"
    pv = "+".join(f"{cf(x)}/(1+{r})^{t}" for t, x in enumerate(cells, 1))
    return f"MAX(0,{pv}+{cf(cells[-1])}*(1+{g})/({r}-{g})/(1+{r})^6)"

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 日月光投控損益（十億新台幣；EPS 新台幣、基本；歸屬母公司）。FY2025A 實績、FY2026 錨點（1Q–3Q26 實績＋4Q26 依指引；三情境相同），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要"
PB = {}; r0 = 70
SCN_T = {"A": "情境 A：顆數基準；委外與份額相對值 1.10／1.15／1.20；全流程 FY2027 US$1.2B；CPO 依 Siliconomics 情境 A 滲透率",
         "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85；相對值 1.05／1.08／1.10；全流程 US$1.0B；CPO 情境 B",
         "C": "情境 C：CY2028–29 顆數 ×0.6；相對值 1.00／0.95／0.92；全流程 US$0.7B；CPO 情境 C"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("atm", "  封測（ATM）", "NT$B", "#,##0.0"), ("leap", "    LEAP 先進封裝", "NT$B", "#,##0.0"),
        ("lsh", "    LEAP ÷ 封測營收", "%", "0.0%"), ("ems", "  EMS", "NT$B", "#,##0.0"), ("agm", "封測毛利率", "%", "0.00%"), ("gm", "毛利率", "%", "0.00%"),
        ("oim", "營業利益率", "%", "0.00%"), ("capr", "資本支出 ÷ 營收", "%", "0.0%"), ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("nim", "  淨利率", "%", "0.00%"),
        ("eps", "EPS", "NT$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{SUMR[('A', 'nim')]}")
IN.cell(INROW["fade_mB"], 5, f"=(PL!G{SUMR[('B', 'nim')]}+PL!G{SUMR[('B', 'nim')]}*0.8)/2")
IN.cell(INROW["fade_mC"], 5, f"=PL!G{SUMR[('C', 'nim')]}*0.8")
for k in ("fade_mA", "fade_mB", "fade_mC"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 錨點").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（NT$B）", "+".join(f"ABS(D{PB[s]['rev']}-a26_rev)" for s in SC)),
               ("  三情境：EPS 差（NT$）", "+".join(f"ABS(D{PB[s]['eps']}-a26_eps)" for s in SC)),
               ("  情境 A：LEAP＋主流 − Q 頁封測（NT$B）", f"D{PB['A']['leap']}+D{PB['A']['ms']}-a26_atm")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 70, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 10; PL.column_dimensions["K"].width = 70
for c in "CDEFGHIJ": PL.column_dimensions[c].width = 13
PL.cell(2, 1, "灰底＝FY2030–FY2032 收斂法（上游只到 2029）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 I 節），非驅動因子推導")

# ---------------- Units：AI 加速器 → LEAP ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — AI 加速器顆數（Siliconomics）→ 日月光 LEAP 先進封裝營收（情境 A）。部署 GW 讀 Siliconomics SIF_DeployGW_A，只作 Recon"
A_ = PB["A"]; cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
GW = SIL.NAMES["SIF_DeployGW_A"]
for key, lab, unit, f, fmt, note in (
        ("units", "AI 加速器顆數（情境 A）", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['units']}", "0.00", "Siliconomics（台積電、Amkor、京元電模型同一份）"),
        ("area", "封裝面積指數（2026＝1）", "x", lambda y: f"=PL!{FC[y]}{A_['area']}", "0.000", "Siliconomics 平台顆數 ÷ CoWoS 每片封裝數"),
        ("rel", "委外與份額相對值", "x", lambda y: f"=PL!{FC[y]}{A_['rel']}", "0.00", "Assumed（情境 A）"),
        ("v", "每顆 LEAP 核心內容", "US$／顆", lambda y: f"=PL!{FC[y]}{A_['v']}", "#,##0", "FY2026 校準；之後隨封裝面積"),
        ("core", "LEAP 核心", "US$B", lambda y: f"=PL!{FC[y]}{A_['core']}", "0.00", "顆數 × 相對值 × 每顆內容"),
        ("fp", "全流程", "US$B", lambda y: f"=PL!{FC[y]}{A_['fp']}", "0.00", "公司 2026 約 0.3、2027 逾 1.0"),
        ("cpo", "CPO", "US$B", lambda y: f"=PL!{FC[y]}{A_['cpo']}", "0.00", "公司 2027 逾 0.5；Siliconomics 滲透率"),
        ("leapu", "LEAP 合計", "US$B", lambda y: f"=PL!{FC[y]}{A_['leapu']}", "0.00", "公司：2026 逾 4.0、2027 翻倍"),
        ("pu", "每顆加速器的日月光 LEAP 營收", "US$／顆", lambda y: f"=PL!{FC[y]}{A_['leapu']}/PL!{FC[y]}{A_['units']}*1000", "#,##0", "含全流程與 CPO"),
        ("gw", "部署 GW（Siliconomics SIF_DeployGW_A）", "GW", lambda y: GW[y - 2026], "0.0", "Siliconomics 推導值（唯讀）"),
        ("pgw", "每 GW 的日月光 LEAP 營收", "US$B／GW", lambda y: f"={cyc(y)}{{leapu}}/{cyc(y)}{{gw}}", "0.000", "對照 Amkor 投資人日：OSAT 先進封裝可服務市場約 US$0.5B／GW")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    UR[key] = ur
    for y in BU:
        fx_ = f(y)
        if isinstance(fx_, str):
            fx_ = fx_.replace("{leapu}", str(UR.get("leapu", 0))).replace("{gw}", str(UR.get("gw", 0)))
        UN[f"{cyc(y)}{ur}"] = fx_; UN[f"{cyc(y)}{ur}"].number_format = fmt
        if key == "gw": UN[f"{cyc(y)}{ur}"].font = BLUE
    UN.cell(ur, 7, note); ur += 1
ur += 1
UN.cell(ur, 1, "讀法：LEAP FY2026 以公司「逾 US$4B」校準，不是公司揭露的拆分；委外與份額相對值是 Assumed，只有相對變化影響 FY2027 起的營收。每顆內容隨封裝面積上升（彈性 0.5），全流程與 CPO 另計；CPO 依 Siliconomics 同字母滲透率縮放。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 52; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 13
UN.column_dimensions["G"].width = 56

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」；目標價＝DCF 與 Comparable 加權（第 0 節）；十年 P/E 規則值（第 3 節，年底落後 P/E）只作參考"
V["A2"] = "循環位置：日月光不是設備股，不套設備市場見頂年；以本模型封測（ATM）營收最高的會計年度為見頂年（Andy 前一串定案；FY(t)＝CY(t)）。封測營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 51
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-30 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "0.00"
for k, y, lab in ((3, 2026, "現價 ÷ FY2026 EPS（錨點）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 B）≈ 0.2 × FY2026＋0.8 × FY2027"); V[f"B{O + 7}"] = f"=0.2*PL!D{SUMR[('B', 'eps')]}+0.8*PL!E{SUMR[('B', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 B）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（Yahoo 奇摩股市，10-27）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis，ADR，10-26）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 3}"] = "FY2026 三情境相同（錨點）"; V[f"E{O + 4}"] = "FY2027 三情境只差委外相對值、全流程與 CPO"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型封測營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, CY0 + 1, ["情境", "封測營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['atm']}:G{PB[s]['atm']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：封測營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；日月光自身封測營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "0.00"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg，Val 第 0.5 節）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
PE_MT = [10.45, 12.36, 7.02, 16.20, 10.35, 6.11, 5.78, 18.78, 21.59, 22.18]
T10 = vr + 2
V.cell(T10, 1, "5. 十年 P/E（2016–2025 年底；單一來源，分位排名用；無負值或 >200 需剔除）").font = BOLD
hdr(V, T10 + 1, ["年底", "採用 P/E", "macrotrends（ADR）", "（第二來源：未找到）"])
RB0 = T10 + 16
for i, y in enumerate(range(2016, 2026)):
    rr = T10 + 2 + i
    V.cell(rr, 1, str(y))
    V.cell(rr, 2, f"=C{rr}").number_format = "0.00"
    V.cell(rr, 3, f"=B{RB0 + 5 + i}").number_format = "0.00"
PE10 = f"$B${T10 + 2}:$B${T10 + 11}"
V.cell(T10 + 12, 1, "年數"); V[f"B{T10 + 12}"] = f"=COUNT({PE10})"
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "macrotrends（ADR 年底 P/E）", "vals": PE_MT, "tag": "Interested-party（資料商；2016–2024 為 12 月底，2025 為 9 月底值，12 月底未列）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位。日月光的落後 P/E 與盈餘反向：盈餘高峰年（2021–2022）約 6x，低谷年（2023–2025）約 19–22x；2016–2017 為合併前的日月光半導體（ADR 代號相同）。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025"
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "0.00"); EPSR[y] = 8 + i
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)" for c in "BCD"], "0.00")
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], "0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底；下限 0")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("Amkor Technology（艾克爾）", "AMKR", 14.02, "2026-10-23"), ("力成（Powertech）", "6239.TW", 20.60, "2026-10-02"),
         ("京元電子（KYEC）", "2449.TW", 26.13, "2026-10-06"), ("矽格（Sigurd）", "6257.TW", 23.30, "2026-10-02"), ("南茂（ChipMOS）", "8150.TW", 22.08, "2026-10-08")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左；與 Amkor、京元電模型同一組數字）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "台美封測同業：Amkor、力成、京元電、矽格、南茂（不放 A 股：長電、通富）")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(2/12*B{EPSR[2027]}+10/12*B{EPSR[2028]})+tp_wB*(2/12*C{EPSR[2027]}+10/12*C{EPSR[2028]})+tp_wC*(2/12*D{EPSR[2027]}+10/12*D{EPSR[2028]})", None, None], "0.00",
     "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "", "", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable（相對倍數調整後，Val 第 0.5 節）；tp_w_dcf 見 In 頁")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "0.00", "賣方平均：未找到；以 BofA 900（2026-10-14，單一券商）代替")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-30 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", "本模型 12 個月後前瞻 EPS ÷ BofA 同口徑前瞻 EPS − 1（In 頁 cons_feps）")
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
tset(46, "band", "價格帶（DCF 不如預期～順利）", ['=TEXT(D23,"0.00")&"～"&TEXT(B23,"0.00")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 5 節）的分位；只作檢查，不進目標價")
V.cell(48, 1, "目標價方法")
V["B48"] = ("DCF：股東現金流＝基本 EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為 FY2023–25 平均約 10.5%（再投資率約 38%），FY2027–32 明確期＋永續；"
            "日月光資本支出約占營收四成，FY2026–27 股東自由現金流為負（PL 頁參考列），本式未扣資本支出高峰，現金流偏高估；但 ROE 用 2023–25 低谷年，FY2026 起 ROE 明顯上升，再投資率偏高（偏低估）。"
            "Comparable：台美封測同業（Amkor、力成、京元電、矽格、南茂）前瞻 P/E 中位數，再乘相對倍數調整係數（Bloomberg 十年前瞻 P/E，日月光相對同業）；不放 A 股。"
            "情境對應：A＝順利（顆數基準）、B＝平平（顆數 ×0.8／0.85）、C＝不如預期（顆數 ×0.6）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16
GR_ROW, RC_ROW = 49, 50
V.cell(GR_ROW, 1, "上檔空間級距")
V.cell(GR_ROW, 2, '=IF(B42>0.5,"大於 50%",IF(B42>0.25,"25%–50%",IF(B42>=0,"0%–25%",IF(B42>=-0.25,"0%～−25%",IF(B42>=-0.5,"−25%～−50%","小於 −50%")))))').number_format = "@"
V.cell(GR_ROW, 5, "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
V.cell(RC_ROW, 1, "投資建議")
V.cell(RC_ROW, 2, '=IF(B42>0.5,"強力買進",IF(B42>0.25,"買進",IF(B42>=-0.25,"中立",IF(B42>=-0.5,"放空","強力放空"))))').number_format = "@"
V.cell(RC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")

# ---------------- Sens（每案三情境重算，對加權目標價；鍵欄） ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（三情境同時重算，再依情境權重求 DCF 與 Comparable 目標價）：對 FY2027、FY2028 EPS（情境 B）與加權目標價的影響。每個案例的三情境推導區塊在本頁下方"
SE["A2"] = "註：情境專屬列（委外相對值、全流程 FY2027、CPO 滲透率、顆數係數）不給敏感度區間，差異見情境 B、C；Comparable 欄已乘相對倍數調整係數（rm_factor），與加權目標價同口徑。"
hdr(SE, 3, ["案例", "鍵", "設定", "FY2027 EPS（B）", "FY2028 EPS（B）", "FY2028 EPS（B）相對基準", "DCF 目標價", "Comparable 目標價", "加權目標價", "加權目標價相對基準"])
CASES = [("AI 部署 ×0.8", "mf", "FY2027 起顆數 ×0.8（三情境）", dict(mf="0.8")), ("AI 部署 ×1.2", "mf", "×1.2", dict(mf="1.2")),
         ("LEAP 每顆內容彈性 0.3", "el", "封裝面積帶動較少", dict(el="el_lo")), ("LEAP 每顆內容彈性 1.0", "el", "與面積等比例", dict(el="el_hi")),
         ("全流程滲透成長 10%", "fp_g", "基準 30%", dict(fpg="fp_g_lo")), ("全流程滲透成長 50%", "fp_g", "", dict(fpg="fp_g_hi")),
         ("CPO 2027 US$0.3B", "cpo27", "基準 0.5（公司逾 0.5）", dict(cpo27="cpo27_lo")), ("CPO 2027 US$0.8B", "cpo27", "", dict(cpo27="cpo27_hi")),
         ("主流封測 FY2027 +5%", "g_ms27", "基準 +15%（公司：與 2026 相近，逾 20%）", dict(gms27="g_ms27_lo")), ("主流封測 FY2027 +22%", "g_ms27", "", dict(gms27="g_ms27_hi")),
         ("主流封測 FY2028–29 +2%", "g_msL", "基準 +7%", dict(gmsL="g_msL_lo")), ("主流封測 FY2028–29 +12%", "g_msL", "", dict(gmsL="g_msL_hi")),
         ("EMS FY2027 0%", "g_ems27", "基準 +10%", dict(gems27="g_ems27_lo")), ("EMS FY2027 +15%", "g_ems27", "", dict(gems27="g_ems27_hi")),
         ("EMS FY2028–29 0%", "g_emsL", "基準 +5%", dict(gemsL="g_emsL_lo")), ("EMS FY2028–29 +8%", "g_emsL", "", dict(gemsL="g_emsL_hi")),
         ("EMS 毛利率 7.5%", "ems_gm", "基準 8.5%", dict(emsgm="ems_gm_lo")), ("EMS 毛利率 9.5%", "ems_gm", "", dict(emsgm="ems_gm_hi")),
         ("匯率 30.5", "fx", "Siliconomics fx_twd 低", dict(fx="fx_lo")), ("匯率 33", "fx", "Siliconomics fx_twd 高", dict(fx="fx_hi")),
         ("資本支出 FY2027 US$11.5B", "capex27", "基準 13.0", dict(capex27="capex27_lo")), ("資本支出 FY2027 US$15B", "capex27", "", dict(capex27="capex27_hi")),
         ("資本支出率 22%", "capi", "FY2028–29；基準 30%", dict(capi="capi_lo")), ("資本支出率 38%", "capi", "", dict(capi="capi_hi")),
         ("折舊年限 6 年", "life", "基準 8 年", dict(life="life_lo")), ("折舊年限 12 年", "life", "", dict(life="life_hi")),
         ("現金流通率 40%", "ft", "基準 50%", dict(ft="ft_lo")), ("現金流通率 60%", "ft", "", dict(ft="ft_hi")),
         ("營業費用彈性 0.3", "phi", "基準 0.6", dict(phi="phi_lo")), ("營業費用彈性 1.0", "phi", "", dict(phi="phi_hi")),
         ("營業外 −8", "nonop_f", "每年；基準 −5", dict(nonop="nonop_f_lo")), ("營業外 −2", "nonop_f", "", dict(nonop="nonop_f_hi")),
         ("稅率 15%", "tax", "基準 17%", dict(tax="tax_lo")), ("稅率 20%", "tax", "", dict(tax="tax_hi"))]
BASE_ENG = 120
SB = {}; brow = BASE_ENG
SE.cell(BASE_ENG - 2, 1, "敏感度「逐案推導區塊」（以下每案三個情境各一區塊；不是案例列）").font = BOLD
for nm, k, st, ov in CASES:
    SB[nm] = {}
    for s in SC:
        SB[nm][s] = engine(SE, brow, f"案例區塊：{nm}（情境 {s}）", scn=s, **ov)
        brow = SB[nm][s]["end"] + 2
YC = [FC[y] for y in YRS6]
def put_case(i, name, key, setting, blocks, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, key); SE.cell(i, 3, setting)
    SE[f"D{i}"] = f"={pre}E{blocks['B']['eps']}"; SE[f"E{i}"] = f"={pre}F{blocks['B']['eps']}"
    SE[f"F{i}"] = f"=E{i}/E$4-1"
    SE[f"G{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_expr([pre + c + str(blocks[s]['eps']) for c in YC])}" for s in SC) + ")*(1+tp_r)"
    feps = "+".join(f"tp_w{s}*(2/12*{pre}E{blocks[s]['eps']}+10/12*{pre}F{blocks[s]['eps']})" for s in SC)
    SE[f"H{i}"] = f"=Val!$B${T['pmed']}*({feps})*rm_factor"
    SE[f"I{i}"] = f"=tp_w_dcf*G{i}+(1-tp_w_dcf)*H{i}"
    SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c, fm in (("D", "0.00"), ("E", "0.00"), ("F", "0.0%"), ("G", "#,##0.0"), ("H", "#,##0.0"), ("I", "#,##0.0"), ("J", "0.0%")):
        SE[f"{c}{i}"].number_format = fm
put_case(4, "基準", "—", "PL 三情境", PB, "PL!")
for n, (nm, k, st, ov) in enumerate(CASES):
    put_case(5 + n, nm, k, st, SB[nm], "")
nn = 5 + len(CASES)
TPS0 = nn
def tp_case(i, name, key, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, key); SE.cell(i, 3, setting)
    SE[f"D{i}"] = "=D4"; SE[f"E{i}"] = "=E4"; SE[f"F{i}"] = "=E{0}/E$4-1".format(i)
    SE[f"G{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_expr(['Val!' + c + str(EPSR[y]) for y in YRS6], r, g)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"H{i}"] = "=H4"
    SE[f"I{i}"] = f"={w}*G{i}+(1-{w})*H{i}"
    SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c, fm in (("D", "0.00"), ("E", "0.00"), ("F", "0.0%"), ("G", "#,##0.0"), ("H", "#,##0.0"), ("I", "#,##0.0"), ("J", "0.0%")):
        SE[f"{c}{i}"].number_format = fm
TPCASES = [("折現率 +1pt", "tp_r", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "tp_r", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "tp_g", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "tp_g", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.3", "tp_w_dcf", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "tp_w_dcf", "Comparable 0.3", dict(w="0.7"))]
for k_, (nm, key, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + k_, nm, key, st, **kw)
TPS_LAST = TPS0 + len(TPCASES) - 1
put_case(TPS_LAST + 1, "情境 B 全權重", "—", "只看情境 B（對照）", {"A": PB["B"], "B": PB["B"], "C": PB["B"]}, "PL!")
put_case(TPS_LAST + 2, "情境 C 全權重", "—", "只看情境 C（對照）", {"A": PB["C"], "B": PB["C"], "C": PB["C"]}, "PL!")
SENS_LAST = TPS_LAST + 2
SE.cell(SENS_LAST + 2, 1, "讀法：顆數情境 B、C 的差異從 FY2028 起；FY2027 只差委外相對值、全流程與 CPO。主流封測約占封測營收七成、EMS 約占營收三成五，所以主流成長率與現金流通率對 EPS 的影響大於 LEAP 單項；折舊年限與資本支出決定 FY2028 起的毛利率。")
SENS_NOTE = SENS_LAST + 2
assert SENS_NOTE < BASE_ENG - 3, SENS_NOTE
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 12; SE.column_dimensions["C"].width = 34
for c in "DEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、公司說法、券商預估與市場口徑。只作事後對照，不為對上數字而調參數。賣方共識未找到，以 BofA（2026-10-14，單一券商，3Q26 財報前）對照"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 4Q26 模型 vs 公司指引（錨點用指引中點，非獨立檢驗）").font = BOLD; rr_ += 1
rec("4Q26 合併營收季增 vs 指引中點 7%（百分點）", "=(g4_rev_lo+g4_rev_hi)/2", lambda s: "=Q!F4/Q!E4-1", "3Q26 新聞稿：6–8%", "", fmt="0.0%", pt=True)
rec("4Q26 合併毛利率 vs 指引中點 23.5%（百分點）", "=(g4_gm_lo+g4_gm_hi)/2", lambda s: "=Q!F15", "3Q26 新聞稿：23–24%", "", fmt="0.0%", pt=True)
rec("4Q26 合併營業利益率 vs 指引中點 15%（百分點）", "=(g4_om_lo+g4_om_hi)/2", lambda s: "=Q!F17/Q!F4", "3Q26 新聞稿：14.5–15.5%", "營業費用季增 8% 校準", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "B. 公司說法（3Q26 法說 2026-10-29）").font = BOLD; rr_ += 1
rec("FY2027 LEAP ÷ FY2026（公司：翻倍）", 2.0, lambda s: f"=PL!E{PB[s]['leapu']}/PL!D{PB[s]['leapu']}", "3Q26 法說：expecting … to double our LEAP revenue", "情境 A 高於翻倍、B 約翻倍", fmt="0.00")
rec("FY2027 全流程（公司：逾 US$1B）", 1.0, pl("fp", 2027), "3Q26 法說：over US$1 billion（保守）", "", fmt="0.00")
rec("FY2027 CPO（公司：逾 US$0.5B）", "=cpo27", pl("cpo", 2027), "3Q26 法說：over US$0.5 billion", "B、C 依 Siliconomics 較低滲透率", fmt="0.00")
rec("FY2027 主流封測年增 vs 公司「與 2026 相近」（FY2026 實測，百分點）", "=Q!G7/Q!B7-1", lambda s: f"=PL!E{PB[s]['ms']}/PL!D{PB[s]['ms']}-1", "3Q26 法說：general market 2027 similar（逾 20%）", "本模型取 15%（偏保守）", fmt="0.0%", pt=True)
rec("FY2027 資本支出 vs FY2026 US$11.5B（公司：更大）", "=capex26", pl("capu", 2027), "3Q26 法說：even larger than this year", "", fmt="0.00")
RC.cell(rr_, 1, "C. 券商（BofA 2026-10-14；3Q26 財報前、單一來源）").font = BOLD; rr_ += 1
rec("FY2026 EPS vs BofA", "=bofa_eps26", pl("eps", 2026), "BofA", "錨點含 3Q26 實績（BofA 為財報前）")
rec("FY2027 EPS vs BofA", "=bofa_eps27", pl("eps", 2027), "BofA", "")
rec("FY2028 EPS vs BofA", "=bofa_eps28", pl("eps", 2028), "BofA", "差距超過 30% 時 README 說明")
rec("FY2028 封測毛利率 vs BofA 約 30%（百分點）", "=bofa_gm28", pl("agm", 2028), "BofA", "", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "D. LEAP 的市場口徑").font = BOLD; rr_ += 1
rec("FY2026 每 GW 的日月光 LEAP 營收（US$B／GW）", None, lambda s: f"=Units!C{UR['pgw']}", "Siliconomics 部署 GW", "對照 Amkor 投資人日：OSAT 先進封裝 US$0.5B／GW（全體 OSAT）", fmt="0.000")
rec("FY2027 每 GW 的日月光 LEAP 營收（US$B／GW）", None, lambda s: f"=Units!D{UR['pgw']}", "同上", "情境 A", fmt="0.000")
rec("FY2028 LEAP ÷ 封測營收（核心問題：三成以上）", 0.30, pl("lsh", 2028), "核心問題門檻", "", fmt="0.0%", pt=True)
RC.column_dimensions["A"].width = 72
for c in "BCDEFGH": RC.column_dimensions[c].width = 11
RC.column_dimensions["I"].width = 40; RC.column_dimensions["J"].width = 60
RECON_LAST = rr_ - 1

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（SEC 6-K 為公司原文；法說內容經逐字稿；券商、媒體與資料商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：3Q26 新聞稿與簡報（SEC 6-K，2026-10-29；3Q26 損益、分部、應用別、4Q26 指引、EBITDA、淨負債）", K3Q),
        ("公司：2Q26 新聞稿與簡報（SEC 6-K，2026-07-30；1Q、2Q26 損益與分部、3Q26 指引、1H26 資本支出）", K2Q),
        ("公司：4Q25 新聞稿與簡報（SEC 6-K，2026-02-05；FY2025、FY2024 損益與分部、LEAP 2025 US$1.6B、資本支出 US$5.5B）", K4Q),
        ("公司：3Q26 法說逐字稿（Motley Fool，2026-10-29；LEAP 逾 US$4B、2027 翻倍、全流程、CPO、資本支出 US$11.5B、稅率）", FOOL3Q),
        ("媒體：數位時代 2Q26 法說（2026-07-30；LEAP 逾 3.5B、全流程 3 億美元、資本支出 10.5B 拆分）", BN2Q),
        ("媒體：風傳媒 2Q26 法說（2026-07-30；營業外項目、3Q26 指引）", STORM2Q),
        ("券商：BofA（科技新報 2026-10-14 摘要；EPS 2026–28、目標價 900、LEAP 占比、資本支出）", BOFA),
        ("價格：富邦證券 e 點通 3711 日線（10-30 收盤 668.00）", FUBON),
        ("資料商：stockanalysis（ADR 價格、前瞻 P/E）", SA),
        ("資料商：macrotrends（ADR 年底 P/E 2016–2025）", MT),
        ("資料商：macrotrends（ROE 2023–2025）", MTROE),
        (f"行業層：{SILV}（加速器顆數、情境係數、CoWoS 每片封裝數、CPO 滲透率、匯率、部署 GW）", "siliconomics/data/" + SIL.FILE),
        ("待補（人工）：賣方共識 EPS 與平均目標價；LEAP 季度金額（公司未揭露）；2.5D 委外比例與份額的量化；十年 P/E 第二來源；FY2016–2023 淨利率", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
OBS = [
    "目標價 NT$504（DCF 337、Comparable 671 各半），比現價 668（2026-10-30）低 24.6%，級距 0%～−25%，投資建議中立。Comparable 單獨為 −10.2%。",
    "核心問題一（LEAP 三成以上後的倍數）：LEAP 占封測營收 FY2026 22.8% → FY2027 35.9%／33.9%／31.8%（A／B／C），現價 ÷ FY2027 EPS 21.6／23.0／24.6 倍，已接近 Bloomberg 2026 窗口前瞻 P/E 24.1 倍、約為十年中位 12.6 倍的兩倍；Comparable 等於 22.1 倍 × 相對倍數係數 1.12＝24.7 倍，若倍數回到十年中位，Comparable 目標價約減半。與 BofA 目標價 900 的差距主要在 EPS（本模型 12 個月後前瞻 EPS 27.2 比 BofA 同口徑 47.2 低 42%），倍數因素只 −14%。",
    "核心問題二（資本支出約占營收四成的報酬）：FY2026 資本支出 US$11.5B（營收 42%）、FY2027 假設 US$13.0B，折舊由 FY2026 NT$97B 升到 FY2028 NT$166B（+71%），營業利益增量 ÷ FY2026 起累計資本支出在 FY2028 只有 12.4%／8.4%／4.7%（A／B／C），股東自由現金流到 FY2029 都為負。加權上檔空間 −24.6% 距放空門檻僅 0.4 個百分點（現金流通率 40% 或折舊年限 6 年即跨線）；下一個檢驗點是 4Q26 財報（2027-02 上旬）的封測毛利率能否達指引 31.6–32.6%。",
]
lines = [(f"日月光投控（3711.TW；ADR ASX）個股模型 {VER}（AI 半導體 Project：封測次產業）　{DATE}", True),
 (f"版本：{VER}（{DATE}：初版）。全球最大封測代工（OSAT）；封測（ATM：封裝、測試、材料）約占營收六成五、EMS（環旭）約三成五。樣板 amkr/；行業變數讀 {SILV}。", False),
 ("模型命題：封測（ATM）＝主流封測＋LEAP 先進封裝；LEAP＝AI 加速器顆數（Siliconomics）× 委外與份額相對值 × 每顆內容（隨封裝面積）＋全流程（CoWoS full process）＋CPO（Siliconomics 同字母滲透率）。公司稱 LEAP 2026 逾 US$4B、2027 翻倍；全流程 2026 約 US$0.3B → 2027 逾 US$1B；CPO 2027 逾 US$0.5B（Interested-party）。EMS 用簡單成長率。", False),
 ("核心問題：LEAP 占封測營收升到三成以上後，前瞻 P/E 應停在歷史中位（約 12.6 倍）與 2026 窗口（約 24.1 倍）之間的哪裡；資本支出約占營收 42% 的報酬能否撐住毛利率與 EPS。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)。預測期 FY2027–FY2029（三年）由驅動因子推導，FY2030–32 為收斂法。金額：十億新台幣；EPS：新台幣（基本）。口徑：歸屬母公司淨利（IFRS）。現價 668.00（2026-10-30 收盤）；其他模型仍用 10-22／10-23 收盤。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("LEAP：Siliconomics 加速器顆數 × 情境係數（2028–29）× 委外與份額相對值（Assumed，情境專屬）× 每顆 LEAP 核心內容（FY2026 以公司「逾 US$4B」校準，之後 × 封裝面積指數 ^ 0.5）＋全流程（FY2027 依情境，之後隨顆數 × 1.3）＋CPO（公司 2027 逾 US$0.5B × Siliconomics 同字母 CPO 滲透率 ÷ 情境 A 2027）。", False),
 ("主流封測：FY2027 +15%（公司稱與 2026 相近，逾 20%；本模型保守）、FY2028–29 +7%（Assumed；行業變數待併入 Siliconomics）。EMS：FY2027 +10%、之後 +5%，毛利率 8.5%。", False),
 ("毛利：封測毛利＝FY2026 錨點＋現金流通率 50% × 封測營收增量 − 折舊增量；折舊＝前一年 × (1 − 1／8)＋前一年資本支出 ÷ 8；資本支出 FY2026 US$11.5B、FY2027 US$13.0B、之後營收 × 30%。營業費用彈性 0.6；營業外每年 −5；稅率 17%。", False),
 ("校準：FY2026 ＝ 1Q–3Q26 實績＋4Q26 依公司指引（ATM 季增 12%、ATM 毛利率 +1.5pt、EMS 持平）。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2）＋0.5 × Comparable 目標價（Amkor、力成、京元電、矽格、南茂前瞻 P/E 中位數 × 12 個月後前瞻 EPS × 相對倍數調整係數）。十年 P/E 規則值保留為參考。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司數字取自 SEC 6-K；法說內容經逐字稿；券商、媒體與資料商為 Interested-party。", False),
 ("", False),
 ("已知限制：(1) LEAP 季度與分項金額公司未揭露，FY2026 以全年說法校準、季度拆分用 BofA 占比；(2) 委外與份額相對值是 Assumed；(3) 主流封測與 EMS 只用成長率；(4) 折舊全數歸封測、混合年限 8 年；(5) DCF 未扣資本支出高峰（偏高估）但 ROE 用低谷年（偏低估）；(6) 賣方共識未找到，以 BofA 單一券商對照；(7) 十年 P/E 只有一個來源；(8) 下次更新：4Q26 財報（2027-02 上旬）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170
_n0 = RD.max_row
RD.cell(_n0 + 2, 1, f'="投資建議："&Val!B{RC_ROW}&"（依上檔空間固定規則）"').font = BOLD
for _i, _t in enumerate(OBS, 1): RD.cell(_n0 + 2 + _i, 1, f"重點觀察 {_i}：" + _t)
OBS_ROW0 = _n0 + 3

# ---------------- 分部原則第 3 條：核對列 ----------------
sys.path.insert(0, os.path.join(ROOT, "tools", "segment_check"))
from addrow import add_check
CHK_SEG = add_check(Q, seg_rows=[5, 8, 9], total_row=4, cols="BCDEFG", rnd=6, note="分部原則第 3 條：官方分部＝封測（ATM）、EMS（分部表口徑），沖銷及其他為合併與分部合計的差額；分部宣告見 segments.json")
Q.cell(CHK_SEG + 2, 1, "檢查：LEAP＋主流 − 封測（命題子列合計）").font = BOLD
for c in "BCDEFG":
    Q[f"{c}{CHK_SEG + 2}"] = f"=ROUND({c}6+{c}7-{c}5,6)"; Q[f"{c}{CHK_SEG + 2}"].number_format = "0.000"
for c in "BCDEFG":
    if Q[f"{c}9"].value and c == "G": pass

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
import relative_multiple as RMX  # noqa: E402
RMX.apply(wb, "ase")
for _row in RD.iter_rows():
    for _c in _row:
        if isinstance(_c.value, str) and _c.value.startswith("2026-10-29：加權目標價中的 Comparable"):
            _c.value = ("相對倍數調整（Val 第 0.5 節）：加權目標價中的 Comparable 使用相對倍數調整後的值（同業前瞻 P/E 中位數 × 本公司 10 年相對同業的前瞻 P/E 比，Bloomberg；"
                        "Andy 2026-10-27「照建議做」）；「Comparable 目標價（市場主流同業法）」不調整。重點觀察 1 的第一句為公式，其餘重點觀察的數字以調整後撰寫。")
sys.path.insert(0, os.path.join(ROOT, "tools", "in_layout")); import apply_in_layout  # noqa: E402
apply_in_layout.apply(wb, VER, DATE)
wb.save(sys.argv[1])
if os.environ.get("ASE_ROWMAP"):
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW, "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE,
               "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE, "UN_NOTE": UN_NOTE, "T": T, "T10": T10, "TPS0": TPS0, "PEER_LAST": PEER_LAST,
               "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0, "OBS_ROW0": OBS_ROW0, "CHK_SEG": CHK_SEG, "README_N": len(lines)},
              open(os.environ["ASE_ROWMAP"], "w"), ensure_ascii=False, indent=1)
