# Alphabet v0.3 成品核對：公式錯誤 0、Q 頁實績核對 0、PL 分部核對（三情境 × 全部年度）0、Seg 核對 0、敏感度未過期、固定列標籤存在。
# 用法：python3 alphabet/check_v03.py <重算後的 xlsx>；失敗時非零結束。
import sys
import openpyxl

X = sys.argv[1]
wb = openpyxl.load_workbook(X, data_only=True)
fail = []
ERR = ("#REF", "#DIV", "#VALUE", "#NAME", "#N/A", "#NUM", "Err:")
n_err = 0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(ERR):
                n_err += 1
                if n_err <= 10: fail.append(f"公式錯誤 {ws.title}!{c.coordinate} {c.value}")
print(f"公式錯誤：{n_err} 格")


def rows_with(ws, prefix):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith(prefix):
            yield r


def nums(ws, r):
    return [c.value for c in ws[r][1:] if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


n_chk = 0
for sh in ("Q", "Seg", "PL", "CF", "Val"):
    ws = wb[sh]
    for r in rows_with(ws, "檢查："):
        if "應為 1" in str(ws.cell(r, 1).value): continue
        vs = nums(ws, r)
        n_chk += 1
        if any(abs(v) > 1e-6 for v in vs):
            fail.append(f"{sh}!A{r}「{ws.cell(r, 1).value}」有非零值 {[round(v, 6) for v in vs if abs(v) > 1e-6][:5]}")
print(f"核對列：{n_chk} 列（Q、Seg、PL、CF、Val；含三情境）")
cf = wb["CF"]
V = wb["Val"]
need = ["目標價（DCF 與 Comparable 加權）", "上檔空間：加權目標價", "投資建議", "上檔空間級距", "DCF 目標價（情境加權）", "Comparable 目標價（市場主流同業法）",
        "價格帶（DCF 不如預期～順利）", "現價（2026-10-26 收盤）", "市場共識目標價"]
labs = {V.cell(r, 1).value: r for r in range(1, V.max_row + 1)}
for n in need:
    if n not in labs: fail.append(f"Val 缺列「{n}」")
if "收支關卡：自由現金流扣股利（五年合計）" not in {cf.cell(r, 1).value for r in range(1, cf.max_row + 1)}:
    fail.append("CF 缺「收支關卡：自由現金流扣股利（五年合計）」")
se = wb["Sens"]
for r in range(1, se.max_row + 1):
    if str(se.cell(r, 1).value or "").startswith("敏感度重算時的基準值"):
        d = se.cell(r, 5).value
        if not isinstance(d, (int, float)) or abs(d) > 1e-6:
            fail.append(f"Sens 敏感度已過期或缺：基準差 {d}")
tp = V.cell(labs["目標價（DCF 與 Comparable 加權）"], 2).value
up = V.cell(labs["上檔空間：加權目標價"], 2).value
rec = V.cell(labs["投資建議"], 2).value
print(f"加權目標價 {tp:.2f}；上檔空間 {up:+.1%}；投資建議 {rec}")
if fail:
    print("失敗：")
    for f in fail: print("  -", f)
    sys.exit(1)
print("全部通過")
