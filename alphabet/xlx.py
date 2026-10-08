import sys, os, json, shutil, subprocess
ROOT = os.path.dirname(os.path.abspath(__file__))  # 暫存檔與重算腳本皆相對於 repo 根目錄
from openpyxl import load_workbook
src, sel, out = sys.argv[1], int(sys.argv[2]), sys.argv[3]
ak = int(sys.argv[4]) if len(sys.argv) > 4 else None  # 選填：EV/EBITDA 錨定年度
os.makedirs(os.path.join(ROOT,'out'),exist_ok=True); p=os.path.join(ROOT,'out',f'xl17_{sel}_{ak or 1}.xlsx'); shutil.copy(src,p)
wb=load_workbook(p); ws=wb['輸入與假設']
for r in range(1,ws.max_row+1):
    v=ws.cell(row=r,column=1).value
    if v and str(v).startswith('情境選擇'): ws.cell(row=r,column=3,value=sel)
    if ak and v and str(v).startswith('EV/EBITDA 錨定年度'): ws.cell(row=r,column=3,value=ak)
wb.save(p); p0=p.replace('.xlsx','_0.xlsx'); wb.save(p0)
r=subprocess.run(['python3',os.path.join(ROOT,'scripts','recalc.py'),p,'120'],capture_output=True,text=True)
if r.returncode: sys.exit('recalc 失敗：'+r.stdout+r.stderr)
# v0.1b：LibreOffice 計算模擬運算表（MULTIPLE.OPERATIONS）時，會把運算表代入其他情境時算出的中間值留在部分儲存格（只在 LibreOffice 發生；
# 例：積極情境 FY30「模型期總營收」殘留基準情境的值）。第二輪：把運算表輸出改為第一輪算出的常數、移除運算表後再重算一次，其餘儲存格全部重新計算。
wv=load_workbook(p,data_only=True)['輸入與假設']; DT={}
for rr in range(1,wv.max_row+1):
    n=wv.cell(row=rr,column=1).value
    if n and str(n).startswith('情境區間運算表｜') and '公式列' not in str(n): DT[rr]=(wv.cell(row=rr,column=4).value,wv.cell(row=rr,column=5).value)
if DT:
    p2=p.replace('.xlsx','_b.xlsx'); wb=load_workbook(p0); ws=wb['輸入與假設']
    for rr,(a,b) in DT.items(): ws.cell(row=rr,column=4,value=a); ws.cell(row=rr,column=5,value=b)
    wb.save(p2)
    r=subprocess.run(['python3',os.path.join(ROOT,'scripts','recalc.py'),p2,'120'],capture_output=True,text=True)
    if r.returncode: sys.exit('recalc（第二輪）失敗：'+r.stdout+r.stderr)
    p=p2
wb=load_workbook(p,data_only=True); o={}
for sh in wb.sheetnames:
    ws=wb[sh]
    for r in range(1,ws.max_row+1):
        n=ws.cell(row=r,column=1).value
        if not n: continue
        if sh=='季度追蹤': o[f"{sh}|{n}"]=[ws.cell(row=r,column=3+i).value for i in range(15)]; continue  # v4.4：季度層（6 季或焦點季 15 欄；含文字）
        if sh=='輸入與假設' and str(n).startswith(('季度', '期間指引')): o[f"{sh}|{n}"]=[ws.cell(row=r,column=3+i).value for i in range(6)]; continue  # v4.4：J 區
        if sh=='評價_可比公司': v=[ws.cell(row=r,column=2+i).value for i in range(8)]
        elif sh=='檢查_連動' and str(n).startswith('關聯方｜'): v=[ws.cell(row=r,column=3+i).value for i in range(5)]  # MAG v0.1b：關聯方並排（C–G 欄）
        elif sh=='檢查_連動': v=[ws.cell(row=r,column=2).value]
        else: v=[ws.cell(row=r,column=3+i).value for i in range(5)]
        if sh=='摘要' or (sh=='輸入與假設' and str(n).startswith('共識｜')): o[f"{sh}|{n}"]=v; continue  # v4.3：一頁摘要與市場共識 I 區（含文字列、擷取日期、標記）
        if any(isinstance(x,(int,float)) for x in v): o[f"{sh}|{n}"]=v
        elif sh=='評價_DCF與目標價' and (n.startswith('目標價區間｜') or n=='DCF 權重說明'): o[f"{sh}|{n}"]=[v[0] or '']  # v4.1：文字列（判斷句等）逐字比對
json.dump(o,open(out,'w'),ensure_ascii=False); print('ok',sel,len(o))
