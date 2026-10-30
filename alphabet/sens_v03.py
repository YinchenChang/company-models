# 敏感度與拆解：一次改一個 In 頁輸入（低／高值），整本以 LibreOffice 重算，讀 Val 頁加權目標價，寫成 JSON 供 build_alphabet.py --sens 使用。
# 用法：python3 alphabet/sens_v03.py <工作資料夾> <輸出 sens.json>（環境變數 JOBS＝平行數，預設 4）
import json, os, subprocess, sys, shutil, tempfile
from concurrent.futures import ThreadPoolExecutor
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WD, OUT = sys.argv[1], sys.argv[2]
os.makedirs(WD, exist_ok=True)
BUILD = os.path.join(HERE, "build_alphabet.py")


Q3CF = os.path.join(HERE, "data", "q3cf_v03_2026q3.json")


def build(name, sets):
    raw = os.path.join(WD, f"{name}.xlsx")
    cmd = [sys.executable, BUILD, raw] + sum((["--q3cf", Q3CF] if k == "__q3cf" else ["--set", f"{k}={v}"] for k, v in sets), [])
    env = dict(os.environ); env["ROWMAP"] = os.path.join(WD, f"{name}.rowmap.json")
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, env=env)
    return raw


def recalc(raw):
    out = os.path.join(WD, "out"); os.makedirs(out, exist_ok=True)
    dst = os.path.join(out, os.path.basename(raw))
    for _ in range(3):                                   # v0.3.3：LibreOffice 偶爾沒有輸出檔 → 重試
        prof = tempfile.mkdtemp(prefix="lo_googl_")
        env = dict(os.environ); env["LC_ALL"] = "C.UTF-8"
        subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", out, raw], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=env, timeout=600)
        shutil.rmtree(prof, ignore_errors=True)
        if os.path.exists(dst): break
    return dst


def tp_of(xlsx, rowmap):
    rm = json.load(open(rowmap, encoding="utf-8"))
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    cf = wb["CF"]
    return {"tp": wb["Val"].cell(rm["VR"]["tp"], 2).value, "gateB": cf.cell(rm["GATE"], 3).value,
            "cx27B": cf[f"H{rm['CR']['B']['cx']}"].value}


def run(name, sets):
    raw = build(name, sets)
    return tp_of(recalc(raw), os.path.join(WD, f"{name}.rowmap.json"))


base_raw = build("base", [])
rm = json.load(open(os.path.join(WD, "base.rowmap.json"), encoding="utf-8"))
IN = openpyxl.load_workbook(base_raw)["In"]
row = rm["IN_ROW"]


def lohi(key, which):
    return IN.cell(row[key], 4 if which == "lo" else 6).value


jobs = {"base|base": []}
PACE0 = json.load(open(os.path.join(HERE, "data", "pace_v031.json"), encoding="utf-8"))
for key, side, val, _r in rm["SE_ROWS"]:
    if key == "c10_win":                                 # C10 下半年視窗
        jobs[f"{key}|{side}"] = [("pace", PACE0["pace_h2"])]
        continue
    if key == "c10_g":                                   # C10 指引區間兩端（指引與速度同改）
        g = 195 if side == "低" else 205
        jobs[f"{key}|{side}"] = [("capex_g", g), ("pace", PACE0[f"pace_g{g}"])]
        continue
    if key == "tpu_fo":                                  # TPU 收入改按產出比（資本支出仍用持有比）
        jobs[f"{key}|{side}"] = [("tpu_f", IN.cell(row["tpu_out"], 5).value)]
        continue
    if key == "tpu_out":                                 # v0.3.3：收入改按產出比的低／高端（tpu_out 本身不進模型公式）
        jobs[f"{key}|{side}"] = [("tpu_f", lohi("tpu_out", val))]
        continue
    if key == "he_all":                                  # v0.3.3：IF_HoldEcon 四世代同時取低／高（原則第 9 節）
        jobs[f"{key}|{side}"] = [(k, lohi(k, val)) for k in ("he_h", "he_b2", "he_b3", "he_vr")]
        continue
    v = lohi(key, val) if val in ("lo", "hi") else val
    sets = [(key, v)]
    if key == "tpu_f":
        sets.append(("tpu_cf", lohi("tpu_cf", val)))
    jobs[f"{key}|{side}"] = sets
# v0.3 → v0.3.1 變動拆解鏈（Recon 頁）：S4＝本版但 (f) 其他改回 v0.3；S3＝S4 再改回舊 TPU 口徑（pace 依舊口徑重解）；
# S2＝S3 再改回 v0.3 的 pace 與伺服器占比；S1＝S2 再以 v0.3 的 2026Q3 模型值取代實績（q3cf_v03.py）
PACE = json.load(open(os.path.join(HERE, "data", "pace_v031.json"), encoding="utf-8"))
CF3 = json.load(open(Q3CF, encoding="utf-8"))
S4 = [("rf", 0.05), ("px_val", 316.80)]
S3 = S4 + [("tpu_f", 1), ("tpu_cf", 1), ("pace", PACE["pace_tpu1"])]
S2 = S4 + [("tpu_f", 1), ("tpu_cf", 1), ("pace", 4159.67), ("cip_it", 0.6)]
S1 = S2 + [("__q3cf", 1)] + [tuple(x) for x in CF3["sets"]]
for k, v in (("S1", S1), ("S2", S2), ("S3", S3), ("S4", S4)):
    jobs[f"chain|{k}"] = v
res = {}
with ThreadPoolExecutor(max_workers=int(os.environ.get("JOBS", "4"))) as ex:
    futs = {k: ex.submit(run, f"c{i:02d}", sets) for i, (k, sets) in enumerate(jobs.items())}
    for k, f in futs.items():
        o = f.result()
        res[k] = {**o, "set": jobs[k]}
        print(k, jobs[k], round(o["tp"], 2), round(o["gateB"], 1), round(o["cx27B"], 1), flush=True)
json.dump(res, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
