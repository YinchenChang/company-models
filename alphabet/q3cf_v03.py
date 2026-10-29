# 變動拆解 (a)／(b) 的反事實：v0.3 對 2026Q3 的模型值（v0.3 只有 2H26 欄 → 依 2025 年同期季節性〔Q3 ÷ (Q3＋Q4)〕拆出 Q3），
# 以及 v0.3 隱含的 2026-09-30 資產負債表輸入。寫 alphabet/data/q3cf_v03_2026q3.json：
#   facts：{事實 id: 值}（build_alphabet.py --q3cf 讀取，取代 2026Q3 實績）；sets：In 頁覆寫（sens_v03.py 的 chain|S1）；compare：新季實績 vs v0.3 模型值（報告用）。
# 用法：python3 alphabet/q3cf_v03.py <v0.3 成品 xlsx>
import json, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
X = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "dist", "_待刪除", "20261027_Alphabet個股模型_v0.3.xlsx")
wb = openpyxl.load_workbook(X, data_only=True)
F0 = {it["id"]: it for g in json.load(open(os.path.join(HERE, "data", "alphabet_facts_20261008.json"), encoding="utf-8"))["groups"].values() for it in g}
F3 = {it["id"]: it for it in json.load(open(os.path.join(HERE, "data", "alphabet_facts_2026q3.json"), encoding="utf-8"))["items"]}


def block(ws, title):
    """情境 B 區塊：回傳 {A 欄標籤: 列號}（自標題列起到下一個情境標題）。"""
    rows = {}; on = False
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "")
        if a.startswith("情境 B"): on = True; continue
        if on and a.startswith("情境 C"): break
        if on and a and a not in rows: rows[a] = r
    return rows


PL, CF = wb["PL"], wb["CF"]
pb, cb = block(PL, "情境 B"), block(CF, "情境 B")
def pl(lab, col="F"): return PL[f"{col}{pb[lab]}"].value
def cf(lab, col="F"): return CF[f"{col}{cb[lab]}"].value
def sea(fid):                                   # 2025 年 Q3 ÷ (Q3＋Q4)
    a, b = F0[f"{fid}.2025q3"]["value"], F0[f"{fid}.2025q4"]["value"]
    return a / (a + b)

out = {}
LINES = [("search", "Google Search & other", "is.rev.search"), ("yt", "YouTube ads", "is.rev.yt"), ("network", "Google Network", "is.rev.network"),
         ("subs", "Google subscriptions, platforms, and devices", "is.rev.subs"), ("cloud", "Google Cloud", "is.rev.cloud")]
for k, lab, fid in LINES:
    out[f"{fid}.2026q3"] = pl(lab) * sea(fid) * 1000
out["is.rev.gs.2026q3"] = sum(out[f"is.rev.{k}.2026q3"] for k in ("search", "yt", "network", "subs"))
out["is.rev.ob.2026q3"] = pl("Other Bets") / 2 * 1000
out["is.rev.hedge.2026q3"] = 0.0
out["is.rev.total.2026q3"] = out["is.rev.gs.2026q3"] + out["is.rev.cloud.2026q3"] + out["is.rev.ob.2026q3"]
out["is.ebit.gs.2026q3"] = pl("Google Services 營業利益") * sea("is.rev.gs") * 1000
out["is.ebit.cloud.2026q3"] = pl("Google Cloud 營業利益") * sea("is.rev.cloud") * 1000
out["is.ebit.ob.2026q3"] = pl("Other Bets 營業利益") / 2 * 1000
abc_lab = [l for l in pb if l.startswith("Alphabet-level activities（避險營收")][0]
out["is.ebit.abc.2026q3"] = pl(abc_lab) / 2 * 1000
out["is.ebit.total.2026q3"] = sum(out[f"is.ebit.{k}.2026q3"] for k in ("gs", "cloud", "ob", "abc"))
half = lambda v: v / 2 * 1000
out["cf.dep.2026q3"] = half(pl("折舊合計（基礎＋新增 AI＋新增非 AI）"))
out["cf.sbc.2026q3"] = half(pl("股權報酬（已含在分部營業利益內）"))
out["cf.cfo.2026q3"] = half(cf("營運現金流"))
out["cf.capex.2026q3"] = -half(cf("資本支出合計"))
out["cf.fcf.2026q3"] = out["cf.cfo.2026q3"] + out["cf.capex.2026q3"]
out["cf.buyback.2026q3"] = 0.0
out["cf.div.2026q3"] = -half(cf("普通股股利"))
ni = half(pl("淨利")); out["is.ni.2026q3"] = ni
out["is.eqGain.2026q3"] = 0.0
# 成本費用：維持實績比例（不影響目標價；只讓 Q 頁有值）
for fid in ("is.cogs", "is.rd", "is.sm", "is.ga", "is.tac"):
    out[f"{fid}.2026q3"] = F3[f"{fid}.2026q3"]["value"] * out["is.rev.total.2026q3"] / F3["is.rev.total.2026q3"]["value"]
out["is.oie.2026q3"] = half(pl("業外收益合計")); out["is.pretax.2026q3"] = out["is.ebit.total.2026q3"] + out["is.oie.2026q3"]
out["is.tax.2026q3"] = out["is.pretax.2026q3"] - ni
she1 = PL[f"E{pb['期末稀釋股數（十億股）']}"].value
out["sh.dilW.2026q3"] = she1 * 1.0025 * 1000
out["is.epsDil.2026q3"] = ni / out["sh.dilW.2026q3"]
# v0.3 隱含的 2026-09-30 輸入（Q3＝2H26 的一半；2026-08 票據 25.0 全數在 Q3）
c0 = 155.411
cash_q3 = c0 + (cf("自由現金流（營運現金 − 資本支出 − 融資租賃本金）") - cf("普通股股利") - cf("特別股股利") - cf("策略投資（Anthropic 出資；1H26 含 Wiz、Intersect 收購）") - cf("債務到期償還")) / 2 + 25.0
sets = [("cash_0", round(cash_q3, 4)), ("debt_0", round(101.085 - 1.0 + 25.0, 4)), ("sh_out", 12.230), ("sh_basic_q2", 12.151),
        ("sh_dil_q2", round(she1 * 1.0025 - (12.230 - 12.151), 6)), ("fines", 17.356), ("spacex", 94.1),
        ("cip_0", round(122.814 + cf("資本支出合計") / 2, 4)), ("da_base", 28.416), ("capex_1h", round(80.598 + cf("資本支出合計") / 2, 4)),
        ("abc_rate", 0.0483), ("sbc_r", 0.0631)]
cmp = {}
for fid, lab in (("is.rev.search", "Google Search & other"), ("is.rev.yt", "YouTube ads"), ("is.rev.network", "Google Network"), ("is.rev.subs", "Subscriptions, platforms, and devices"),
                 ("is.rev.gs", "Google Services"), ("is.rev.cloud", "Google Cloud"), ("is.rev.ob", "Other Bets"), ("is.rev.total", "合併營收"),
                 ("is.ebit.gs", "Google Services 營業利益"), ("is.ebit.cloud", "Google Cloud 營業利益"), ("is.ebit.ob", "Other Bets 營業利益"),
                 ("is.ebit.abc", "Alphabet-level"), ("is.ebit.total", "合併營業利益"), ("cf.dep", "折舊"), ("cf.capex", "資本支出（負值）"), ("cf.cfo", "營運現金流")):
    a, m = F3[f"{fid}.2026q3"]["value"] / 1000, out[f"{fid}.2026q3"] / 1000
    cmp[lab] = {"actual": round(a, 3), "v03": round(m, 3), "diff": round(a - m, 3), "pct": round(a / m - 1, 4) if m else None}
cmp["Cloud AI（v0.3 模型 2H26 一半，估計）"] = {"v03": round(pl("Cloud AI（命題子列）") / 2, 3)}
J = {"_readme": "v0.3 對 2026Q3 的模型值（2H26 依 2025 年 Q3 ÷ (Q3＋Q4) 拆分；Other Bets、Alphabet 層級、折舊、資本支出、現金流取 2H26 的一半）與 v0.3 隱含的 2026-09-30 輸入；"
                 "供 v0.3.1 變動拆解 (a)／(b) 使用（alphabet/q3cf_v03.py 由 dist/_待刪除/20261027_Alphabet個股模型_v0.3.xlsx 產生）",
     "facts": out, "sets": sets, "compare": cmp, "cash_q3": cash_q3}
json.dump(J, open(os.path.join(HERE, "data", "q3cf_v03_2026q3.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for k, v in cmp.items(): print(k, v)
print(sets)
