# 離線開啟檢查（已決定事項 11）：交付的 HTML 必須是單一檔案、完全離線、以 file:// 雙擊即可開啟。
# 做法：把 HTML 單獨複製到空的暫存資料夾（旁邊沒有任何檔案），以 Playwright 在阻斷網路的狀態下用 file:// 開啟，檢查：
#   (a) 網路請求數＝0：除了 HTML 本身以外，任何請求（http/https/ws、其他本機檔案）都算失敗，且一律攔截不放行；
#   (b) console 無錯誤、頁面無未捕捉例外；
#   (c) 一頁摘要與主要分頁的關鍵數字正常顯示：數字取自同版 Excel（重算後的快取值），依 A 欄列名稱查找，
#       依格式轉成畫面字串後須出現在該分頁的可見文字中；各分頁也不得出現 NaN／undefined／Infinity。
# 用法：python3 scripts/check_offline.py 檔案.html 同版.xlsx
# 5b 之後改為依具名範圍取值（CHECKS 的「分頁、列名稱」改為名稱），檢查項目不變。
import sys, os, re, shutil, tempfile
from pathlib import Path
from urllib.parse import unquote
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from first_screen import JS_VIEWPORT_TEXT

# （畫面分頁路徑, Excel 工作表, A 欄列名稱, 欄號, 格式）；格式：usd1＝$47.9、usd2＝$90.13、n1＝151.7、n2＝3.66、text＝原字串
CHECKS = [
    ('總結', '摘要', '結論｜評等', 3, 'text'),
    ('總結', '摘要', '結論｜點位（加權目標價）', 3, 'usd1'),
    ('總結', '摘要', '結論｜現價', 3, 'usd2'),
    ('總結', '摘要', '結論｜情境區間（下緣／上緣）', 3, 'usd1'),
    ('總結', '摘要', '結論｜情境區間（下緣／上緣）', 4, 'usd1'),
    ('資金模型', '檢查_連動', '模型期毛 CapEx（«VMD» 起）', 2, 'n1'),
    ('資金模型', '檢查_連動', '模型期營運缺口（«VMD» 起）', 2, 'n1'),
    ('損益與評價', '摘要', '差異｜營收｜模型', 4, 'n1'),              # FY27 營收（損益簡表）
    ('損益與評價/市場共識', '輸入與假設', '共識｜目標價｜平均', 3, 'usd2'),
    ('資金模型/各期收支/季度追蹤', '季度追蹤', '焦點｜營收', 3, 'n2'),
]
# v4.5（5a-1）：第一屏的日期與期間文字（只比對 1440×900 視窗內、捲動到頂端時可見的文字；已決定事項 9 第 3、4 輪）
# 格式：text＝原字串；md＝2026-06-30 → 6/30
CHECKS_FIRST = [
    ('總結', '輸入與假設', '資料更新日', 3, 'text'),
    ('總結', '輸入與假設', '模型期間', 3, 'text'),
    ('資金模型', '輸入與假設', '模型期間', 3, 'text'),
    ('損益與評價', '輸入與假設', '現價日（收盤）', 3, 'text'),
    ('損益與評價', '輸入與假設', '評價日（已申報季度的季末）', 3, 'md'),
]
BAD_TEXT = ('NaN', 'undefined', 'Infinity')


def xl_value(wb, sheet, label, col, unit=None):
    # 列名稱中的期間佔位符（«VMD» 等，v4.5）以任意文字比對，新舊兩版的 Excel 都能找到
    parts = re.split(r'(«[A-Z0-9]+»)', label)   # 偶數位置為原文、奇數位置為佔位符
    pat = re.compile('^' + ''.join('.+?' if i % 2 else re.escape(x) for i, x in enumerate(parts)) + '$') if len(parts) > 1 else None
    ws = wb[sheet]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if (a == label or (pat and isinstance(a, str) and pat.match(a))) and ws.cell(r, col).value is not None:
            if unit is not None: unit.append(ws.cell(r, 2).value)  # WhiteFiber v0.1c：B 欄單位（US$m 時畫面金額 1 位小數）
            return ws.cell(r, col).value
    raise KeyError(f'Excel 找不到：{sheet}｜{label}（第 {col} 欄）')


def fmt(v, how):
    if how == 'text': return str(v)
    if how == 'md': y, m, d = str(v)[:10].split('-'); return f'{int(m)}/{int(d)}'
    s = {'usd1': f'${v:,.1f}', 'usd2': f'${v:,.2f}', 'n1': f'{v:,.1f}', 'n2': f'{v:,.2f}'}[how]
    return s


def main(html, xlsx):
    wb = load_workbook(xlsx, data_only=True)
    expect, first = {}, {}
    for tab, sh, lab, col, how in CHECKS:
        u = []; v = xl_value(wb, sh, lab, col, u)
        if how == 'n2' and u and u[0] == 'US$m': how = 'n1'  # WhiteFiber v0.1c：US$m 金額畫面最多 1 位小數（DUQ）
        expect.setdefault(tab, []).append((f'{sh}｜{lab}', fmt(v, how)))
    for tab, sh, lab, col, how in CHECKS_FIRST:
        try: v = xl_value(wb, sh, lab, col)
        except KeyError:
            if sh == '輸入與假設': continue  # v4.4 以前的 Excel 沒有「期間與日期」列（v4.5 新增）：略過
            raise
        first.setdefault(tab, []).append((f'{sh}｜{lab}', fmt(v, how)))
        expect.setdefault(tab, [])
    tmp = tempfile.mkdtemp(prefix='offline_')
    page_path = Path(tmp) / Path(html).name          # 單獨一個檔案，旁邊沒有任何其他檔案
    shutil.copy(html, page_path)
    main_url = page_path.as_uri()
    reqs, errs, fails = [], [], []
    with sync_playwright() as p:
        try: b = p.chromium.launch()
        except Exception:  # 與 crawl.py 相同：預裝 Chromium 版本可能與 playwright 套件不符
            b = p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH', '/opt/pw-browsers/chromium'))
        ctx = b.new_context(offline=True, viewport={'width': 1440, 'height': 900})
        same = lambda u: unquote(u) == unquote(main_url)
        # 只放行 HTML 本身；其他請求（網路、旁邊的本機檔案）全部記錄並攔截。data:／blob: 是內嵌內容，不算請求
        ctx.route('**/*', lambda route: route.continue_() if same(route.request.url) else route.abort())
        ctx.on('request', lambda r: None if same(r.url) or r.url.startswith(('data:', 'blob:', 'about:')) else reqs.append(r.url))
        pg = ctx.new_page()
        pg.on('console', lambda m: errs.append('console：' + m.text) if m.type == 'error' else None)
        pg.on('pageerror', lambda e: errs.append('例外：' + str(e)))
        pg.goto(main_url); pg.wait_for_timeout(2500)

        def click(t):
            pg.locator('button:visible').filter(has_text=re.compile('^\\s*' + re.escape(t) + '\\s*$')).first.click(timeout=3000)
            pg.wait_for_timeout(500)
        for tab, items in expect.items():
            try:
                for i, t in enumerate(tab.split('/')):
                    if i == 0: pg.get_by_role('button', name=t, exact=True).click(); pg.wait_for_timeout(700)
                    else: click(t)
            except Exception as e:
                fails.append(f'{tab}：無法切換到此分頁（{str(e).splitlines()[0]}）'); continue
            pg.evaluate('window.scrollTo(0, 0)'); pg.wait_for_timeout(150)
            vp = pg.evaluate(JS_VIEWPORT_TEXT)  # 第一屏（視窗內）文字
            for key, s in first.get(tab, []):
                if s not in vp: fails.append(f'{tab}：第一屏找不到 {key}＝{s}')
            text = pg.inner_text('body')
            if len(text) < 500: fails.append(f'{tab}：畫面文字只有 {len(text)} 字（未正常顯示）')
            for w in BAD_TEXT:
                if w in text: fails.append(f'{tab}：畫面出現「{w}」')
            for key, s in items:
                if s not in text: fails.append(f'{tab}：找不到 {key}＝{s}')
        b.close()
    shutil.rmtree(tmp, ignore_errors=True)
    n = sum(len(v) for v in expect.values()) + sum(len(v) for v in first.values())
    print(f'(a) 網路請求：{len(reqs)} 個' + ('' if not reqs else '：' + '、'.join(reqs[:5])))
    print(f'(b) console 錯誤與例外：{len(errs)} 個' + ('' if not errs else '：' + '；'.join(e[:120] for e in errs[:5])))
    print(f'(c) 關鍵數字：{len(expect)} 個分頁、{n} 項，失敗 {len(fails)} 項' + ('' if not fails else '\n   ' + '\n   '.join(fails)))
    ok = not reqs and not errs and not fails
    print('離線開啟檢查：' + ('通過' if ok else '不通過'))
    return 0 if ok else 1


if __name__ == '__main__':
    if len(sys.argv) != 3: sys.exit('用法：python3 scripts/check_offline.py 檔案.html 同版.xlsx')
    sys.exit(main(sys.argv[1], sys.argv[2]))
