# AMD（AMD）收支與估值模型 v0.2.2（AI 半導體 Project：運算層，GPU 側非 NVIDIA 份額）
# 用法：python3 amd/build_amd.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年（FY2026＝CY2026；AMD 的會計年度約在 12 月最後一個週六結束，不做平移）。金額：十億美元（US$B）；EPS 為 non-GAAP（公司口徑）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "nvidia"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))   # 本資料夾的 tp_v2.py 優先
import market as MEM
import cycle_position as CP
import valuation_rule as VR
import tp_v2 as TP
import accel_share as AS          # 加速器金額份額的可重用函式（NVIDIA 模型建立；AMD 為 GPU 側非 NVIDIA 份額）

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
PNAME = {k: nm for k, nm, *_ in MEM.UNITS}
PLAT = [k for k, *_ in MEM.UNITS]

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); NV_F = only("nvidia/dist/*.xlsx"); BRG_F = only("markets/chip_bridge/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)

ws_ = openpyxl.load_workbook(TSMC_F, data_only=True)["In"]
TS = {}   # 台積電模型 In 頁：顆數、2029 成長、情境顆數係數
for r in ws_.iter_rows(min_row=4):
    k = r[0].value
    if k and (k.startswith("um2") or k.startswith("g29_") or k.startswith("u_")):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
for k, nm, u_, gb, tag, src in MEM.UNITS:     # 與 markets/memory 的 UNITS 必須一致（同一份）
    for i, y in enumerate((26, 27, 28)):
        assert abs(TS[f"u_{k}_{y}"][0] - u_[i]) < 1e-9, (k, y)
for s_ in SC:
    assert abs(TS[f"um28_{s_}"][0] - MEM.SCEN_UNITS[s_][0]) < 1e-9 and abs(TS[f"um29_{s_}"][0] - MEM.SCEN_UNITS[s_][1]) < 1e-9
GB_AMD = [g for kk, nm, u_, g, t, sr in MEM.UNITS if kk == "amd"][0]

wsn = openpyxl.load_workbook(NV_F, data_only=True)["In"]
NV = {}   # NVIDIA 模型 In 頁：各平台每顆價（份額計算同一套價格）、HBM 價格路徑（skhynix 成品轉抄）
for r in wsn.iter_rows(min_row=4):
    k = r[0].value
    if k and (k.startswith("pa_") or k.startswith("pm_") or k.startswith("hbm_")):
        NV[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(NV_F)}｜In!D{r[0].row}:F{r[0].row}", r[7].value)
for k in ("pa_amd", "pa_tpu", "pa_trn", "pa_oth", "pm_bw", "pm_rb", "pm_ru", "pm_hp") + tuple(f"hbm_{s}_{y}" for s in SC for y in (26, 27, 28, 29)):
    assert k in NV and NV[k][0][1] is not None, k

wbb = openpyxl.load_workbook(BRG_F, data_only=True)["In"]
assert wbb["A7"].value == "gpus" and isinstance(wbb["H7"].value, (int, float)), wbb["A7"].value
GPW = (wbb["H7"].value, f"{rel(BRG_F)}｜In!H7（VR200 每 GW GPU 數，Tokenomics IF_GPUsPerGW）")

PEAK, PEAK_SRC = CP.peak_years(ROOT)   # 見頂年度（日曆年）：設備市場模型
assert tuple(CP.classify(2027, PEAK[s_]) for s_ in SC) == ("中段", "中段", "見頂"), PEAK

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；每顆價格：千美元（US$K）；顆數：百萬顆；股數：十億股。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未單獨使用。公司原文＝AMD 8-K 新聞稿與簡報（sec.gov 已逐頁抓取）；法說逐字稿只經摘要（二手）"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
E8K = "https://www.sec.gov/Archives/edgar/data/2488/"
Q226 = E8K + "000000248826000121/q22026991.htm"; Q226S = E8K + "000000248826000121/amdq22026earningsslidesf.htm"
Q126 = E8K + "000000248826000072/q12026991.htm"; Q425 = E8K + "000000248826000014/q42025991.htm"

sec("A. 加速器顆數（markets/memory 與 HBM、台積電、網通、NVIDIA 模型同一份；情境係數讀台積電模型成品，不另訂）")
for k, nm, u_, gb, tag, src in MEM.UNITS:
    for i, y in enumerate((26, 27, 28)):
        inp(f"u_{k}_{y}", f"{nm} 顆數 20{y}", "百萬顆", None, u_[i], None, tag, "markets/memory/market.py UNITS：" + src if i == 0 else "同上", f"台積電模型 {TS[f'u_{k}_{y}'][1]}（值相同，建置時檢查）" if i == 0 else "同上")
    inp(f"g29_{k}", f"{nm} 2029 顆數成長", "%", None, TS[f"g29_{k}"][0], None, "Assumed", TS[f"g29_{k}"][1], "markets/memory 無 2029；沿用台積電模型")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
for i, y in enumerate((26, 27, 28)):
    inp(f"gb_amd_{y}", f"AMD Instinct 每顆 HBM 20{y}", "GB", None, GB_AMD[i], None, "Interested-party", "markets/memory/market.py UNITS（MI355X 288；MI450 432）" if i == 0 else "同上", "世代售價係數的 Analogy 區間以容量比計（見 D 節）" if i == 0 else "同上")

sec("B. 各平台每顆售價（金額份額用；與 NVIDIA 模型同一套價格，讀 NVIDIA 成品 In 頁；模組層級）")
for k in ("pm_bw", "pm_rb", "pm_ru", "pm_hp", "pa_amd", "pa_tpu", "pa_trn", "pa_oth"):
    lo, b, hi = NV[k][0]
    inp(k, f"每顆售價：{({'pm_bw': PNAME['bw'], 'pm_rb': PNAME['rb'], 'pm_ru': PNAME['ru'], 'pm_hp': PNAME['hp'], 'pa_amd': PNAME['amd'], 'pa_tpu': PNAME['tpu'], 'pa_trn': PNAME['trn'], 'pa_oth': PNAME['oth']})[k]}", "US$K", lo, b, hi, "Derived", NV[k][1], (NV[k][2] or "")[:160] or "NVIDIA 模型輸入（標記見該模型）")

sec("C. HBM 價格路徑（情境別；讀 NVIDIA 成品 In 頁，原始來源 skhynix 成品 Mkt 頁；唯讀）")
for s in SC:
    for y in (26, 27, 28, 29):
        inp(f"hbm_{s}_{y}", f"HBM 混合單價 20{y}（{s}）", "US$/GB", None, NV[f"hbm_{s}_{y}"][0][1], None, "Derived", NV[f"hbm_{s}_{y}"][1], "markets/memory：2026 BofA $14.6/GB；2027 +80% 三情境相同；2028–29 依情境" if (s == "A" and y == 26) else "同上")

sec("D. AMD 資料中心 GPU（Instinct）：公司不拆分 GPU 與 CPU，拆分比例為 Assumed")
inp("gpu25", "FY2025 資料中心 GPU 營收（歷史對照；不進預測）", "US$B", 5.9, 6.5, 7.5, "Assumed", "公司不揭露。Next Platform 估 Q4 2025 Instinct 2.65（含中國 MI308 0.39）；Wells Fargo 估 Q4 約 2.5–2.6（二手）；Q2 2025 有 MI308 出口管制費用 $0.8B（公司）", "2024 Instinct 逾 $5B（公司 2025-01 法說，二手）。各季拆分：未找到")
inp("gs_q1", "Q1 2026 GPU 占資料中心營收", "%", 0.38, 0.42, 0.46, "Assumed", "公司不揭露；Q4 2025 Next Platform 估 GPU 約占 49%（含中國）；Q1 2026 中國 MI308 約 $0.1B（公司指引）", "第二來源：未找到")
inp("gs_q2", "Q2 2026 GPU 占資料中心營收", "%", 0.39, 0.43, 0.47, "Assumed", "公司：資料中心 AI 營收年增逾一倍（Q2 2025 基期受 MI308 管制壓低）；伺服器 CPU 雲端與企業皆年增逾 70%（法說摘要，二手：MarketBeat）", "第二來源：未找到")
inp("ng_q3", "Q3 2026 資料中心非 GPU（CPU 等）季增", "%", 0.04, 0.08, 0.12, "Assumed", "公司：下半年伺服器營收年增逾 80%、2026 伺服器 CPU 供給吃緊（法說摘要，二手）", "—")
inp("ng_q4", "Q4 2026 資料中心非 GPU 季增", "%", 0.02, 0.06, 0.10, "Assumed", "同上；Venice（Zen 6）2026 下半年開始部署（公司）", "—")
inp("gm_gpu", "資料中心 GPU 毛利率（FY2026；non-GAAP 口徑）", "%", 0.40, 0.50, 0.55, "Assumed", "公司：資料中心 AI 毛利率低於公司平均（Q2 2026 法說摘要，二手）；公司平均 FY2026 錨點約 56%", "第二來源：未找到。非 GPU 毛利率由此在 FY2026 校準（Derived，PL 各區塊）")
inp("abs_mem", "HBM 漲價由 AMD 吸收的比例（其餘轉嫁到售價）：FY2026–27 三情境共用；FY2028 起改讀 K 節 abs_mem_A／_B／_C；低／高為 Sens 用", "x", 0.0, 0.5, 1.0, "Assumed", "比照 NVIDIA 模型（Assumed 0–1，基準 0.5）", "第二來源：未找到。AMD 對 HBM4 合約價轉嫁的公開說法：未找到")
inp("uf", "AMD 顆數係數（Sens 用；FY2027 起乘以 AMD 市場層顆數）", "x", 0.8, 1.0, 1.2, "Assumed", "無：OpenAI、Meta 各 6 GW、Anthropic 2 GW、Oracle 5 萬顆等大單的時點不確定（見 H 節）", "±20%")
for y, gbx in ((27, GB_AMD[1]), (28, GB_AMD[2]), (29, GB_AMD[2])):
    inp(f"pg_{y}", f"世代售價係數 20{y}（MI450／Helios 每顆售價 ÷ FY2026 認列售價）", "x", 1.0, 1.0, round(gbx / GB_AMD[0], 6), "Analogy／Assumed", "基準 1.0＝k_rec 延續（比照 NVIDIA 模型：FY2026 校準的水準沿用）；高＝售價隨 HBM 容量等比例（Analogy：432 ÷ 288）", "MI455X 官方價：未找到；部落格 $15–50K（NVIDIA 模型 In 頁 pa_amd 說明）")

sec("E. 其他業務（簡化成長率；FY2027 起）")
inp("g_cpu27", "資料中心非 GPU（EPYC CPU 等）年增 FY2027", "%", 0.30, 0.50, 0.70, "Assumed", "公司：2027 伺服器 CPU 年增逾 70%（Q2 2026 法說摘要，二手；只作 Recon，高＝公司說法）", "伺服器 CPU 市場 2030 年約 $220B、年增逾 50%（公司，Interested-party）")
inp("g_cpu28", "資料中心非 GPU 年增 FY2028", "%", 0.10, 0.25, 0.40, "Assumed", "無", "—")
inp("g_cpu29", "資料中心非 GPU 年增 FY2029", "%", 0.05, 0.15, 0.25, "Assumed", "無", "—")
inp("g_cli", "Client（PC 處理器）年增（FY2027 起）", "%", -0.05, 0.03, 0.10, "Assumed", "公司：2026 下半年 PC 市場轉弱、季節性低於上半年（二手）；Q2 2026 Client 年增 23%（公司）", "第二來源：未找到")
inp("g_gam", "Gaming（主機半客製＋顯示卡）年增（FY2027 起）", "%", -0.25, -0.10, 0.10, "Assumed", "公司：2026 半客製營收大幅雙位數下滑（Q4 2025 法說，二手）；Q2 2026 Gaming 年減 31%（公司）", "下一代主機時程：未找到")
inp("g_emb", "Embedded 年增（FY2027 起）", "%", 0.0, 0.10, 0.15, "Assumed", "公司：Q2 2026 年增 19%、Q3 強勁雙位數季增；今年設計案逾 $18B（二手）", "—")

sec("F. FY2026 錨點假設（Q1、Q2 實績＋Q3 指引＋Q4 假設）")
inp("q3_rev", "Q3 2026 營收指引（±$0.3B）", "US$B", 12.7, 13.0, 13.3, "Interested-party", Q226 + "：約 $13B ±$300M", "MarketBeat、platformaeronaut 摘要（二手）")
inp("q3_gm", "Q3 2026 non-GAAP 毛利率指引", "%", None, 0.56, None, "Interested-party", Q226 + "：約 56%", "Q2 2026 簡報（同一公司文件）")
inp("q3_opx", "Q3 2026 non-GAAP 營業費用指引", "US$B", None, 3.65, None, "Interested-party", Q226S + "：約 $3.65B", "platformaeronaut 摘要（二手）")
inp("q3_oth", "Q3 2026 non-GAAP 其他收益（扣利息費用）指引", "US$B", None, 0.055, None, "Interested-party", Q226S + "：約 $55M", "—")
inp("tx_26", "FY2026 non-GAAP 有效稅率", "%", None, 0.13, None, "Interested-party", Q226S + "：13%", "Q1–Q2 實績推得約 13.0%（Q 頁）")
inp("q3_sh", "Q3 2026 稀釋股數指引", "十億股", None, 1.661, None, "Interested-party", Q226S + "：約 1,661M", "—")
inp("q3_cli", "Q3 2026 Client 季增（假設）", "%", 0.0, 0.02, 0.05, "Assumed", "公司：Client 與 Gaming 合計小幅季減；Client 略增、Gaming 強勁雙位數下滑（法說摘要，二手）", "—")
inp("q3_gam", "Q3 2026 Gaming 季增（假設）", "%", -0.30, -0.20, -0.10, "Assumed", "同上", "—")
inp("q3_emb", "Q3 2026 Embedded 季增（假設）", "%", 0.08, 0.12, 0.16, "Assumed", "公司：強勁雙位數季增（二手）", "—")
inp("q4_dc", "Q4 2026 資料中心季增（假設）", "%", 0.10, 0.25, 0.40, "Assumed", "公司：Helios 第三季底開始出貨、第四季放量（Q2 2026 法說摘要，二手）；MI450 第四季有顯著量（Q4 2025 法說）", "共識 FY2026 營收 $51.0B 隱含 Q4 約 $16.2B（只作對照，不校準）")
inp("q4_cli", "Q4 2026 Client 季增（假設）", "%", -0.10, -0.05, 0.02, "Assumed", "公司：下半年季節性低於上半年（二手）", "—")
inp("q4_gam", "Q4 2026 Gaming 季增（假設）", "%", -0.20, -0.10, 0.10, "Assumed", "無", "—")
inp("q4_emb", "Q4 2026 Embedded 季增（假設）", "%", 0.0, 0.03, 0.08, "Assumed", "無", "—")
inp("q4_gm", "Q4 2026 non-GAAP 毛利率（假設）", "%", 0.54, 0.555, 0.57, "Assumed", "公司：資料中心 AI 毛利率低於平均，MI450 放量季組合偏低（二手）", "高＝Q3 指引附近")
inp("q4_opx_g", "Q4 2026 營業費用季增（假設）", "%", 0.02, 0.05, 0.08, "Assumed", "Q1→Q2 +7.9%、Q2→Q3 指引 +7.5%（公司）；下半年費用成長落後營收（公司，二手）", "—")
inp("q4_sh", "Q4 2026 稀釋股數季變動", "%", 0.0, 0.002, 0.01, "Assumed", "Q1→Q2 +0.5%（公司）", "認股權證的稀釋另計（G 節）")

sec("G. 營業費用、稅、股數、認股權證（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.5, "=Q!B36", 0.9, "Derived", "FY2026 錨點營業費用年增 ÷ 營收年增（Q 頁）", "低／高 Assumed；公司：費用成長低於營收成長（二手）")
inp("g_opx_min", "營業費用年增下限", "%", 0.0, 0.05, 0.10, "Assumed", "無", "營收下滑時費用僵固")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準為 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 ±2 個百分點")
inp("tx_f", "non-GAAP 有效稅率（FY2027 起）", "%", 0.12, 0.13, 0.15, "Interested-party", "公司 FY2026 13%；FY2025 推得 13.0%（Q 頁）", "—")
inp("oth_f", "non-GAAP 其他收益淨額（扣利息；FY2027 起每年）", "US$B", 0.0, 0.2, 0.4, "Assumed", "Q3 2026 指引 $55M／季（公司）；Q1–Q2 2026 推得 0.06、0.08（Q 頁）", "—")
inp("sh_g", "稀釋股數年變動（不含認股權證；FY2027 起）", "%", -0.005, 0.01, 0.015, "Derived", "FY2025 1.636 → Q2 2026 1.659（年化約 +1.4%，公司）；2026 上半年買回僅 $0.22B（公司 10-Q）", "低／高 Assumed")
inp("w_tot", "客戶認股權證上限：OpenAI 160M ＋ Meta 160M 股", "十億股", None, 0.32, None, "Verified", "公司 2025-10-06 OpenAI 新聞稿；2026-02-24 Meta 新聞稿（8-K EX-99.1）", "依 GW 出貨里程碑與股價門檻分批歸屬；各批股數：未找到")
inp("w_gw", "認股權證對應的 GW 合計（OpenAI 6＋Meta 6）", "GW", None, 12, None, "Verified", "同上", "—")
inp("gw_26", "已歸屬 GW 累計：2026 年底", "GW", 0.0, 0.0, 1.0, "Assumed", "首批隨首個 1 GW 歸屬；兩案首 1 GW 皆 2026 下半年開始出貨（公司）", "歸屬時點：未找到")
inp("gw_27", "已歸屬 GW 累計：2027 年底", "GW", 1.0, 2.0, 3.0, "Assumed", "假設兩案各完成 1 GW", "股價門檻（OpenAI 最後一批 $600，媒體報導，二手）未模型化")
inp("gw_28", "已歸屬 GW 累計：2028 年底（情境 B、C 乘顆數係數）", "GW", 2.0, 4.0, 6.0, "Assumed", "無", "—")
inp("gw_29", "已歸屬 GW 累計：2029 年底（情境 B、C 乘顆數係數）", "GW", 4.0, 6.0, 9.0, "Assumed", "無", "—")
inp("wf", "認股權證歸屬速度係數（Sens 用）", "x", 0.0, 1.0, 1.5, "Assumed", "0＝不歸屬（股價門檻或里程碑未達）", "—")

sec("H. 公司指引、大單與共識（只作 Recon 對帳，不作校準）")
inp("g_dc27", "公司：2027 資料中心營收年增（逾一倍）", "%", None, 1.0, None, "Interested-party", "Q2 2026 法說（MarketBeat、platformaeronaut 摘要，二手）", "逐字稿未取得")
inp("g_srv27", "公司：2027 伺服器 CPU 年增（逾 70%）", "%", None, 0.70, None, "Interested-party", "同上", "下半年 2026 逾 80%")
inp("lt_eps", "公司：策略期間 EPS 逾 $20（Q2 2026 稱將顯著超過）", "US$", None, 20, None, "Interested-party", "2025-11 分析師日（Q4 2025 法說重述，Motley Fool 逐字稿）", "營收年複合成長逾 35%、資料中心逾 60%")
inp("openai_gw", "OpenAI：6 GW（首 1 GW MI450 於 2026 下半年）；CFO：數百億美元營收", "GW", None, 6, None, "Interested-party", "AMD 2025-10-06 新聞稿（ir.amd.com）", "OpenAI 為交易對手（立場：取得認股權證）")
inp("meta_gw", "Meta：最多 6 GW（首 1 GW 客製 MI450，2026 下半年出貨）", "GW", None, 6, None, "Interested-party", "AMD 2026-02-24 新聞稿（8-K EX-99.1）", "—")
inp("anth_gw", "Anthropic：最多 2 GW MI450（首 1 GW 自 2027 上半年）", "GW", None, 2, None, "Interested-party", Q226 + "；MarketBeat 摘要", "—")
inp("orcl_u", "Oracle：首批 5 萬顆 MI450（2026 第三季起，2027 起擴大）", "百萬顆", None, 0.05, None, "Interested-party", "Oracle 2025-10-14 新聞稿（oracle.com）", "Oracle 為交易對手")
inp("gpw", "每 GW GPU 數（VR200 NVL72；Helios 每架 72 顆的 Analogy）", "顆", None, GPW[0], None, "Analogy", GPW[1], "MI450 每 GW 顆數：未找到")
inp("p355", "MI355X 每顆售價（HSBC）", "US$K", 20, 23, 25, "Interested-party", "HSBC 約 $23K（自 $25K 下修；finviz 轉述）", "其他報導 $20–25K")
inp("cons_rev26", "共識 FY2026 營收（S&P Global，48 位）", "US$B", 48.3, 51.0, 52.7, "Interested-party", "stockanalysis.com（2026-10-23）", "另一來源 tickflow（2026-10-24）47.15，可能較舊")
inp("cons_eps26", "共識 FY2026 non-GAAP EPS", "US$", 7.00, 7.59, 8.20, "Interested-party", "stockanalysis.com（2026-10-23）", "tickflow 6.81")
inp("cons_rev27", "共識 FY2027 營收", "US$B", None, 89.17, None, "Interested-party", "stockanalysis.com「Revenue Next Year」（2026-10-23）", "tickflow 四季合計 67.1（各季分析師數不一，可能較舊）")
inp("cons_eps27", "共識 FY2027 non-GAAP EPS", "US$", 11.23, 15.76, None, "Interested-party", "stockanalysis.com「EPS Next Year」（2026-10-23）", "低＝tickflow 四季合計 11.23（本模型算術）；兩來源差 40%")
inp("cons_eps28", "共識 FY2028 EPS（約 $20）", "US$", 14.41, 20.0, None, "Interested-party", "ts2.tech 2026-07-30 轉述 Motley Fool（二手）", "低＝tickflow 四季合計 14.41")

sec("I. 價格與估值")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 516.79, None, "Interested-party", "stockanalysis.com 歷史股價與統計頁（2026-10-23 收盤 516.79）", "digrin.com AMD 價格頁現價 516.79（2026-10-25 頁面，即 10-23 收盤）；heygotrade 每日報導的收盤序列與 stockanalysis 一致（10-19 538.05、10-14 570.82）")
inp("r_disc", "折現率（把 FY2027 年底隱含價折回現在；參考用）", "%", 0.08, 0.10, 0.12, "Assumed", "無", "FY2027 年底（2027-12）距 2026-10-23 約 1.19 年")
for s in SC:
    inp(f"pk_{s}", f"情境 {s}：設備市場見頂年度（日曆年；循環位置規則用）", "年", None, PEAK[s], None, "Derived", PEAK_SRC[s], "共用規則（tools/cycle_position.py classify）：AMD 會計年度＝日曆年，不平移" if s == "A" else "同上")
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 0.3／0.5／0.2）", "三權重合計應為 1")
sec("J. 目標價參數（v0.2 新增：DCF 與 Comparable 加權、FY2029 後收斂法；黃底＝Andy 可改）")
inp("tp_r", "DCF 折現率（股東權益成本）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25：美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（FY2032 後永續；PL 延伸期 B 情境的收斂目標）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單 2026-10-25）", "Sens 頁 ±1pt（只改 DCF 再投資率與永續成長，不改 PL 收斂路徑）")
inp("tp_gd", "延伸期 A／C 情境長期成長相對 tp_g 的差（A＋、C−）", "pt", None, 0.02, None, "Assumed", "工作單 2026-10-25：A＝tp_g＋2pt、C＝tp_g−2pt", "—")
inp("tp_roe", "ROE（再投資率＝tp_g ÷ tp_roe）：同口徑近三年平均未找到，用工作單預設 0.20", "%", None, 0.20, None, "Assumed", "工作單：近三年平均 ROE，找不到用 0.20。GAAP ROE（valueinvesting.io）2023 1.53%、2024 2.85%、2025 6.88%，平均 3.75%，受 Xilinx 併購攤銷壓低，與本模型 non-GAAP EPS 不同口徑，不採", "若採 GAAP 3.75%（≤ tp_g），再投資率達上限 0.8，股東現金流只剩 EPS 的 20%；stockanalysis 近 12 月 ROE 10.20%")
inp("tp_w_dcf", "目標價中 DCF 的權重（其餘為 Comparable）", "%", None, 0.5, None, "Assumed", "Andy 可改（Andy 2026-10-25 11:53 定 50／50）", "Sens 頁 0.3／0.7")
inp("nm10_k", "延伸期 C 情境淨利率目標＝FY2029 淨利率 × 本係數（同口徑十年平均未找到）", "x", None, 0.8, None, "Assumed", "工作單：找不到十年平均用最後一年 × 0.8。本模型為 non-GAAP 口徑；同口徑十年淨利率：未找到", "參考：GAAP 十年平均 7.10%（valueinvesting.io 2016–2025：−11.63%、−0.63%、5.20%、5.07%、25.50%、19.24%、5.59%、3.77%、6.36%、12.51%），含 Xilinx 攤銷，口徑不同不採")
inp("cons_tp", "市場共識目標價（賣方平均）", "US$", None, 637.61, None, "Interested-party", "stockanalysis.com（55 位分析師平均，2026-10-25 頁面；現價為 10-23 收盤）", "賣方有目標價誘因；只作對照，不進目標價")
sec("K. 情境一致性（v0.2.1 新增，Andy 2026-10-25 13:13「情境與行業一致」：公司特有變數方向 A 最有利、C 最不利）")
for s, v, why in (("A", 0.2, "順利：加速器需求強，HBM 漲價轉嫁最多"), ("B", 0.5, "平平：與 v0.2 三情境共用值相同"), ("C", 0.7, "不如預期：需求轉弱，AMD 吸收最多以保份額")):
    inp(f"abs_mem_{s}", f"情境 {s}：HBM 漲價由 AMD 吸收的比例（FY2028 起）", "x", None, v, None, "Assumed", why, "v0.2 為三情境 0.5，HBM 價 A 最高使 EPS B＞A（稽核 2026-10-25）；改為 A 0.2／B 0.5／C 0.7（A 先試 0.3：FY2028 EPS 11.07＜B 11.11，因認股權證稀釋 A 最多，故取 0.2）。轉嫁比例的公開資料：未找到")

r = 4
INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（十億美元；EPS 為美元；股數為十億股；non-GAAP 公司口徑）。會計年度＝日曆年。藍字＝輸入；黑字＝公式。實績取自 AMD 8-K 新聞稿（sec.gov）；第二來源＝後續季度新聞稿的比較欄（同一公司不同文件）與 10-Q"
Q["A2"] = "FY2026 錨點＝Q1、Q2 實績＋Q3 指引中點＋Q4 假設（I、J 欄，假設見 In 頁 F 節）。資料中心 GPU／非 GPU 拆分公司不揭露，為 Assumed（In 頁 D 節）。EPS 口徑：non-GAAP（排除股票薪酬、併購無形資產攤銷、投資損益等；公司口徑），共識亦為此口徑"
hdr(Q, 3, ["項目", "Q1 2025", "Q2 2025", "Q3 2025", "Q4 2025", "FY2025 全年", "Q1 2026", "Q2 2026", "Q3 2026（指引中點）", "Q4 2026（假設）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "AMD 8-K 新聞稿：Q1 2026 " + Q126 + "；Q2 2026 " + Q226 + "；Q4 2025（含 Q3 2025、全年）" + Q425
S2 = "Q1、Q2 2025 見 Q1、Q2 2026 稿比較欄；Q2 2025 毛利、利息、其他收益為全年 − 其餘三季（Derived）"
ACT = "BCDEGH"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None, ann26=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(B{r}:E{r})"; Q[f"F{r}"].number_format = fmt
        if ann is not None and not (isinstance(ann, str) and ann.startswith("=")): Q[f"F{r}"].font = BLUE
    if ann26 != "":
        Q[f"K{r}"] = ann26 or f"=SUM(G{r}:J{r})"; Q[f"K{r}"].number_format = fmt
    Q.cell(r, 12, tag); Q.cell(r, 13, s1); Q.cell(r, 14, s2)
def act(vs): return dict(zip(ACT, vs))
qrow(4, "營收", {**act([7.438, 7.685, 9.246, 10.270, 10.253, 11.536]), "I": "=q3_rev", "J": "=J5+J8+J9+J10"}, tag="Verified／Interested-party（Q3 指引）／Derived（Q4＝各段合計）")
qrow(5, "Data Center", {**act([3.674, 3.240, 4.341, 5.380, 5.775, 6.718]), "I": "=I4-I8-I9-I10", "J": "=I5*(1+q4_dc)"}, tag="Verified／Derived（Q3＝指引 − 其他段）／Assumed（Q4）", s2=S2 + "；FY2025 合計 16.635＝公司全年值")
qrow(6, "  其中：資料中心 GPU（Instinct；Assumed 拆分）", {"G": "=G5*gs_q1", "H": "=H5*gs_q2", "I": "=I5-I7", "J": "=J5-J7"}, ann="=gpu25", tag="Assumed／Derived", s1="公司不揭露 GPU 與 CPU 拆分；FY2025 為全年估計（In 頁 gpu25），各季：未找到", s2="Q3、Q4＝資料中心 − 非 GPU")
qrow(7, "  其中：資料中心非 GPU（EPYC CPU、DPU、FPGA 等）", {"G": "=G5-G6", "H": "=H5-H6", "I": "=H7*(1+ng_q3)", "J": "=I7*(1+ng_q4)"}, ann="=F5-F6", tag="Derived／Assumed（Q3、Q4 季增）", s1="資料中心 − GPU", s2="—")
qrow(8, "Client", {**act([2.294, 2.499, 2.750, 3.097, 2.885, 3.062]), "I": "=H8*(1+q3_cli)", "J": "=I8*(1+q4_cli)"}, tag="Verified／Assumed（I、J）")
qrow(9, "Gaming", {**act([0.647, 1.122, 1.298, 0.843, 0.720, 0.779]), "I": "=H9*(1+q3_gam)", "J": "=I9*(1+q4_gam)"}, tag="Verified／Assumed（I、J）")
qrow(10, "Embedded", {**act([0.823, 0.824, 0.857, 0.950, 0.873, 0.977]), "I": "=H10*(1+q3_emb)", "J": "=I10*(1+q4_emb)"}, tag="Verified／Assumed（I、J）")
qrow(11, "檢查：營收 − 各段合計（應為 0）", {c: f"={c}4-{c}5-{c}8-{c}9-{c}10" for c in "BCDEFGHIJK"}, fmt="0.000", ann="", ann26="", tag="—", s1="—", s2="—")
qrow(12, "毛利（non-GAAP）", {**act([3.992, 3.326, 4.992, 5.855, 5.685, 6.488]), "I": "=I4*q3_gm", "J": "=J4*q4_gm"}, tag="Verified（Q2 2025 Derived；Q2 2026＝營業利益＋營業費用）／Interested-party（I）／Assumed（J）", s2="FY2025 合計 18.165＝公司全年值")
qrow(13, "毛利率（non-GAAP）", {c: f"={c}12/{c}4" for c in "BCDEFGHIJK"}, fmt="0.0%", ann="", ann26="", tag="Derived", s1="—", s2="公司：Q2 2025 43%（含 MI308 費用 $0.8B）、FY2025 52%、Q2 2026 56%")
qrow(14, "營業費用（non-GAAP）", {**act([2.213, 2.429, 2.754, 3.001, 3.145, 3.394]), "I": "=q3_opx", "J": "=I14*(1+q4_opx_g)"}, tag="Verified／Interested-party（I）／Assumed（J）")
qrow(15, "營業利益（non-GAAP）", {**act([1.779, 0.897, 2.238, 2.854, 2.540, 3.094]), "I": "=I12-I14", "J": "=J12-J14"}, tag="Verified／Derived（I、J）")
qrow(16, "檢查：毛利 − 營業費用 − 營業利益（實績應為 0）", {c: f"={c}12-{c}14-{c}15" for c in "BCDEGH"}, fmt="0.000", ann="", ann26="", tag="—", s1="—", s2="—")
qrow(17, "利息費用", act([-0.020, -0.038, -0.037, -0.036, -0.037, -0.037]), ann26="", tag="Verified（Q2 2025 Derived）")
qrow(18, "其他收益淨額（GAAP）", act([0.039, 0.098, 0.082, 0.358, 0.165, 0.598]), ann26="", tag="Verified（Q2 2025 Derived）")
qrow(19, "  長期投資（利得）損失（non-GAAP 調節，排除）", act([0.002, -0.061, -0.026, -0.280, -0.066, -0.483]), ann26="", tag="Verified（Q2 2025 Derived）", s2="non-GAAP 調節表「(Gains) losses on long-term investments, net」")
qrow(20, "非營業淨額（non-GAAP）＝利息＋其他收益＋調節", {**{c: f"={c}17+{c}18+{c}19" for c in "BCDEGH"}, "I": "=q3_oth", "J": "=q3_oth"}, tag="Derived／Interested-party（I）／Assumed（J＝Q3 指引）")
qrow(21, "稅前利益（non-GAAP）", {c: f"={c}15+{c}20" for c in "BCDEGHIJ"}, tag="Derived")
qrow(22, "淨利（non-GAAP）", {**act([1.566, 0.781, 1.965, 2.519, 2.265, 2.760]), "I": "=I21*(1-tx_26)", "J": "=J21*(1-tx_26)"}, tag="Verified／Derived（I、J）", s2=S2 + "；FY2025 合計 6.831＝公司全年值")
qrow(23, "有效稅率（non-GAAP；1 − 淨利 ÷ 稅前）", {**{c: f"=1-{c}22/{c}21" for c in "BCDEGH"}, "I": "=tx_26", "J": "=tx_26", "F": "=1-F22/F21", "K": "=1-K22/K21"}, fmt="0.0%", ann="", ann26="", tag="Derived／Interested-party（13%）", s1="—", s2="公司 FY2026 non-GAAP 稅率 13%（Q2 2026 簡報）")
qrow(24, "稀釋股數（十億股）", {**act([1.626, 1.630, 1.641, 1.649, 1.650, 1.659]), "I": "=q3_sh", "J": "=I24*(1+q4_sh)"}, ann=1.636, ann26="=AVERAGE(G24:J24)", tag="Verified／Interested-party（I）／Assumed（J）", s2="Q2 2025 1.630 見 Q2 2026 10-Q；FY2025 1.636（公司全年值）")
qrow(25, "EPS（non-GAAP，模型算）＝淨利 ÷ 股數", {**{c: f"={c}22/{c}24" for c in "BCDEGHIJ"}, "F": "=F22/F24", "K": "=K22/K24"}, fmt="0.000", ann="", ann26="", tag="Derived", s1="—", s2="FY2026 Q3、Q4 為指引與假設推得")
qrow(26, "EPS（non-GAAP 稀釋，公司）", act([0.96, 0.48, 1.20, 1.53, 1.37, 1.66]), fmt="0.00", ann=4.17, ann26="", tag="Verified")
qrow(27, "檢查：模型算 EPS − 公司 EPS（四捨五入差）", {c: f"={c}25-{c}26" for c in "BCDEFGH"}, fmt="0.000", ann="", ann26="", tag="—", s1="—", s2="—")
qrow(28, "EPS（GAAP 稀釋，公司；參考）", act([0.44, 0.54, 0.75, 0.92, 0.84, 1.38]), fmt="0.00", ann=2.65, ann26="", tag="Verified", s2="GAAP 含併購無形資產攤銷（FY2025 $2.25B）、股票薪酬（$1.64B）")
qrow(29, "檢查：FY2025 四季合計 − 公司全年值（營收 34.639、營業利益 7.768、淨利 6.831，應為 0）", {"F": "=(F4-34.639)+(F15-7.768)+(F22-6.831)"}, fmt="0.000", ann="", ann26="", tag="—", s1="—", s2="—")
for nm, rr in (("rev", 4), ("dc", 5), ("gpu", 6), ("dcx", 7), ("cli", 8), ("gam", 9), ("emb", 10), ("gp", 12), ("opx", 14), ("oi", 15), ("oth", 20), ("ni", 22), ("etr", 23), ("sh", 24), ("eps", 25)):
    defname(f"a25_{nm}", f"Q!$F${rr}"); defname(f"a26_{nm}", f"Q!$K${rr}")
Q.cell(31, 1, "FY2026 錨點與共識的距離（參考；校準時未使用共識）").font = BOLD
Q.cell(32, 1, "FY2026 營收 ÷ 共識營收 − 1"); Q["B32"] = "=K4/cons_rev26-1"; Q["B32"].number_format = "0.0%"
Q.cell(33, 1, "FY2026 EPS ÷ 共識 EPS − 1"); Q["B33"] = "=K25/cons_eps26-1"; Q["B33"].number_format = "0.0%"
Q.cell(34, 1, "FY2026 營收年增"); Q["B34"] = "=K4/F4-1"; Q["B34"].number_format = "0.0%"
Q.cell(35, 1, "FY2026 營業費用年增"); Q["B35"] = "=K14/F14-1"; Q["B35"].number_format = "0.0%"
Q.cell(36, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B36"] = "=B35/B34"; Q["B36"].number_format = "0.00"
Q.column_dimensions["A"].width = 62
for c in "BCDEFGHIJK": Q.column_dimensions[c].width = 14
Q.column_dimensions["L"].width = 24; Q.column_dimensions["M"].width = 90; Q.column_dimensions["N"].width = 70
Q.freeze_panes = "B4"

# ---------------- Units：加速器顆數、金額份額（accel_share）、校準 ----------------
CYS = (2026, 2027, 2028, 2029)
cyc = lambda y: chr(ord("C") + y - 2026)          # Units 的日曆年欄：C＝2026 … F＝2029
def uexp(k, y, s):
    if y <= 2027: return f"u_{k}_{y % 100}"
    if y == 2028: return f"u_{k}_28*um28_{s}"
    return f"u_{k}_28*(1+g29_{k})*um29_{s}"
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 加速器顆數（百萬顆；日曆年；與 HBM、台積電、網通、NVIDIA 模型同一份）與金額份額（模組層級每顆價 × 顆數；價格與 NVIDIA 模型相同）。AMD 會計年度＝日曆年"
hdr(UN, 3, ["項目", "單位", "CY2026（FY2026）", "CY2027（FY2027）", "CY2028（FY2028）", "CY2029（FY2029）", "說明"])
PRICE = {"bw": "pm_bw", "rb": "pm_rb", "ru": "pm_ru", "hp": "pm_hp", "amd": "pa_amd", "tpu": "pa_tpu", "trn": "pa_trn", "oth": "pa_oth"}
ycols = {y: cyc(y) for y in CYS}
SHR = {}; ur = 4
for s in SC:
    ttl = {"A": "情境 A：基準", "B": "情境 B：2028 顆數 ×0.8、2029 ×0.85", "C": "情境 C：2028–29 顆數 ×0.6"}[s]
    units = {(p, y): uexp(p, y, s) for p in PLAT for y in CYS}
    ur, SHR[s] = AS.write_share_block(UN, ur, ttl, CYS, ycols, units, PRICE, PLAT, PNAME)
UN.cell(ur, 1, "金額份額的口徑：各平台「模組層級」每顆價（NVIDIA＝GPU 模組價；AMD、ASIC＝每顆晶片價；In 頁 B 節，與 NVIDIA 模型相同），不含系統加價與網通。AMD 資料中心 GPU 營收＝加速器市場金額 × AMD 金額份額 × k_rec × 世代售價係數（PL 頁）。").font = BOLD
ur += 2
CAL = ur
UN.cell(CAL, 1, "校準（只用公司已發生的實績與當年指引：FY2026 錨點的資料中心 GPU 營收；不用共識、不用公司對 2027 的說法）").font = BOLD
UN.cell(CAL + 1, 1, "AMD 認列係數 k_rec ＝ FY2026 錨點 GPU 營收 ÷（AMD 市場層 CY2026 顆數 × 市場層每顆價 pa_amd）"); UN.cell(CAL + 1, 2, "x")
UN[f"C{CAL+1}"] = "=a26_gpu/(u_amd_26*pa_amd)"; UN[f"C{CAL+1}"].number_format = "0.0000"; UN[f"D{CAL+1}"] = "Derived"
UN.cell(CAL + 1, 7, "顆數層與售價層無法分開識別，k_rec 是合併係數；FY2027 起沿用（比照 NVIDIA 模型）。錨點 GPU 本身含 Assumed 拆分（In 頁 D 節）")
defname("k_rec", f"Units!$C${CAL+1}")
UN.cell(CAL + 2, 1, "FY2026 隱含每顆認列售價（k_rec × pa_amd；若 k_rec 全歸售價層）"); UN.cell(CAL + 2, 2, "US$K")
UN[f"C{CAL+2}"] = "=k_rec*pa_amd"; UN[f"C{CAL+2}"].number_format = "0.0"
UN.cell(CAL + 2, 7, "對照 HSBC MI355X 約 $23K（Recon）：k_rec 多數可由 2026 以 MI355X 為主的售價解釋")
defname("p_real26", f"Units!$C${CAL+2}")
UN.cell(CAL + 3, 1, "HBM 成本占 FY2026 每顆認列售價 hs_26 ＝ 每顆 HBM 容量 × 2026 HBM 混合單價 ÷ 每顆認列售價"); UN.cell(CAL + 3, 2, "x")
UN[f"C{CAL+3}"] = "=gb_amd_26*hbm_A_26/1000/p_real26"; UN[f"C{CAL+3}"].number_format = "0.0000"; UN[f"D{CAL+3}"] = "Derived"
UN.cell(CAL + 3, 7, "FY2027 起 HBM 成本占比 ＝ hs_26 × HBM 價格指數（容量增加視為隨世代售價等比例，Analogy）")
defname("hs_26", f"Units!$C${CAL+3}")
UN.cell(CAL + 4, 1, "AMD 認列顆數 FY2026（市場層 × k_rec）"); UN.cell(CAL + 4, 2, "百萬顆"); UN[f"C{CAL+4}"] = "=u_amd_26*k_rec"; UN[f"C{CAL+4}"].number_format = "0.00"
UN.column_dimensions["A"].width = 70; UN.column_dimensions["B"].width = 16
for c in "CDEF": UN.column_dimensions[c].width = 18
UN.column_dimensions["G"].width = 80

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
DEFP = dict(scn="A", uf="1", pg=lambda y: f"pg_{y % 100}", abs="abs_mem", abs_s=None, gmg="gm_gpu", gcpu=lambda y: f"g_cpu{y % 100}", opxa="0", wf="wf")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    absy = lambda fy: P["abs"] if fy <= 2027 else (P["abs_s"] or f"abs_mem_{s}")   # v0.2.1：FY2027 三情境共用 abs_mem；FY2028 起讀 abs_mem_{s}
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.0", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda fy: FC[fy]
    uc = lambda fy: f"Units!{cyc(fy)}"
    sec_("A. 由下而上（C 欄空白；D＝FY2026 … G＝FY2029；FY2026 由校準等於錨點，FY2027 起為驅動因子推導）")
    row("ufac", "AMD 認列係數（k_rec × Sens 顆數係數）", "x", {"D": "=k_rec", **{c_(fy): f"=k_rec*{P['uf']}" for fy in BU[1:]}}, fmt="0.000")
    row("mkt", "加速器市場金額（Units 頁，模組層級）", "US$B", {c_(fy): f"={uc(fy)}{SHR[s]['dtot']}" for fy in BU})
    row("shr", "AMD 金額份額（Units 頁）", "%", {c_(fy): f"={uc(fy)}{SHR[s][('ds', 'amd')]}" for fy in BU}, fmt="0.0%")
    row("uamd", "AMD 市場層顆數", "百萬顆", {c_(fy): f"={uc(fy)}{SHR[s][('u', 'amd')]}" for fy in BU}, fmt="0.00")
    row("urec", "AMD 認列顆數（市場層 × 認列係數）", "百萬顆", {c_(fy): f"={c_(fy)}{R['uamd']}*{c_(fy)}{R['ufac']}" for fy in BU}, fmt="0.00")
    row("pg", "世代售價係數（MI450／Helios）", "x", {"D": "=1", **{c_(fy): "=" + P["pg"](fy) for fy in BU[1:]}}, fmt="0.000")
    row("gbase", "GPU 基準營收 ＝ 市場金額 × AMD 份額 × 認列係數 × 世代係數", "US$B", {c_(fy): f"={c_(fy)}{R['mkt']}*{c_(fy)}{R['shr']}*{c_(fy)}{R['ufac']}*{c_(fy)}{R['pg']}" for fy in BU})
    row("hix", "HBM 價格指數（相對 2026）", "x", {c_(fy): f"=hbm_{s}_{fy % 100}/hbm_{s}_26" for fy in BU}, fmt="0.000")
    row("gpur", f"資料中心 GPU 營收（含 HBM 漲價轉嫁 (1−吸收比例)；吸收比例 FY2026–27 abs_mem、FY2028 起 abs_mem_{s}）", "US$B", {c_(fy): f"={c_(fy)}{R['gbase']}*(1+(1-{absy(fy)})*hs_26*({c_(fy)}{R['hix']}-1))" for fy in BU}, note="v0.2.1：FY2028 起吸收比例分情境（A 0.2 轉嫁最多、B 0.5、C 0.7）；FY2026–27 三情境相同（abs_mem）")
    row("gpuc", "資料中心 GPU 成本", "US$B", {c_(fy): f"={c_(fy)}{R['gbase']}*((1-{P['gmg']})+hs_26*({c_(fy)}{R['hix']}-1))" for fy in BU}, note="基準營收 ×（非 HBM 成本＋HBM 成本占比的增加）")
    row("gpugm", "資料中心 GPU 毛利率", "%", {c_(fy): f"=1-{c_(fy)}{R['gpuc']}/{c_(fy)}{R['gpur']}" for fy in BU}, fmt="0.0%")
    row("cpu", "資料中心非 GPU（EPYC CPU 等）", "US$B", {"D": "=a26_dcx", **{c_(fy): f"={c_(fy - 1)}{r}*(1+{P['gcpu'](fy)})" for fy in BU[1:]}})
    row("cli", "Client", "US$B", {"D": "=a26_cli", **{c_(fy): f"={c_(fy - 1)}{r}*(1+g_cli)" for fy in BU[1:]}})
    row("gam", "Gaming", "US$B", {"D": "=a26_gam", **{c_(fy): f"={c_(fy - 1)}{r}*(1+g_gam)" for fy in BU[1:]}})
    row("emb", "Embedded", "US$B", {"D": "=a26_emb", **{c_(fy): f"={c_(fy - 1)}{r}*(1+g_emb)" for fy in BU[1:]}})
    row("othr", "非 GPU 營收合計", "US$B", {c_(fy): f"={c_(fy)}{R['cpu']}+{c_(fy)}{R['cli']}+{c_(fy)}{R['gam']}+{c_(fy)}{R['emb']}" for fy in BU})
    row("go", "非 GPU 毛利率（FY2026 校準：錨點毛利 − GPU 毛利；FY2027 起沿用）", "%", {"D": f"=(a26_gp-(D{R['gpur']}-D{R['gpuc']}))/D{R['othr']}", **{c_(fy): f"=$D{r}" for fy in BU[1:]}}, fmt="0.0%", note="Derived；組合內（CPU、Client、Gaming、Embedded）毛利率差異不模型化")
    row("othc", "非 GPU 成本", "US$B", {c_(fy): f"=(1-{c_(fy)}{R['go']})*{c_(fy)}{R['othr']}" for fy in BU})
    row("bur", "由下而上營收合計", "US$B", {c_(fy): f"={c_(fy)}{R['gpur']}+{c_(fy)}{R['othr']}" for fy in BU})
    row("buc", "由下而上成本合計", "US$B", {c_(fy): f"={c_(fy)}{R['gpuc']}+{c_(fy)}{R['othc']}" for fy in BU})
    row("bugm", "由下而上毛利率", "%", {c_(fy): f"=1-{c_(fy)}{R['buc']}/{c_(fy)}{R['bur']}" for fy in BU}, fmt="0.0%", note="FY2026 ＝ 錨點毛利率（校準恆等式，非獨立檢驗）")
    sec_("B. 報告值（FY2025A 實績、FY2026 錨點；FY2027 起為由下而上）")
    row("rev", "營收", "US$B", {"C": "=a25_rev", "D": "=a26_rev", **{c_(fy): f"={c_(fy)}{R['bur']}" for fy in BU[1:]}})
    row("dc", "  Data Center", "US$B", {"C": "=a25_dc", "D": "=a26_dc", **{c_(fy): f"={c_(fy)}{R['gpur']}+{c_(fy)}{R['cpu']}" for fy in BU[1:]}})
    row("gpu", "    資料中心 GPU", "US$B", {"C": "=a25_gpu", "D": "=a26_gpu", **{c_(fy): f"={c_(fy)}{R['gpur']}" for fy in BU[1:]}})
    row("dcx", "    資料中心非 GPU（CPU 等）", "US$B", {"C": "=a25_dcx", "D": "=a26_dcx", **{c_(fy): f"={c_(fy)}{R['cpu']}" for fy in BU[1:]}})
    row("cg", "  Client＋Gaming", "US$B", {"C": "=a25_cli+a25_gam", "D": "=a26_cli+a26_gam", **{c_(fy): f"={c_(fy)}{R['cli']}+{c_(fy)}{R['gam']}" for fy in BU[1:]}})
    row("embr", "  Embedded", "US$B", {"C": "=a25_emb", "D": "=a26_emb", **{c_(fy): f"={c_(fy)}{R['emb']}" for fy in BU[1:]}})
    row("g", "營收年增", "%", {c_(fy): f"={c_(fy)}{R['rev']}/{c_(fy - 1)}{R['rev']}-1" for fy in FYS[1:]}, fmt="0.0%")
    row("gp", "毛利（non-GAAP）", "US$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(fy): f"={c_(fy)}{R['rev']}-{c_(fy)}{R['buc']}" for fy in BU[1:]}})
    row("gm", "毛利率", "%", {c_(fy): f"={c_(fy)}{R['gp']}/{c_(fy)}{R['rev']}" for fy in FYS}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(fy): f"={c_(fy - 1)}{r}*(1+MAX(g_opx_min,phi*{c_(fy)}{R['g']}))" for fy in BU[1:]}}, note="前一年 ×（1＋MAX(下限, 彈性 × 營收年增)）")
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(fy): f"={c_(fy)}{R['opb']}+{P['opxa']}*{c_(fy)}{R['rev']}" for fy in BU[1:]}})
    row("oi", "營業利益", "US$B", {c_(fy): f"={c_(fy)}{R['gp']}-{c_(fy)}{R['opx']}" for fy in FYS})
    row("oim", "營業利益率", "%", {c_(fy): f"={c_(fy)}{R['oi']}/{c_(fy)}{R['rev']}" for fy in FYS}, fmt="0.0%")
    row("oth", "非營業淨額（non-GAAP）", "US$B", {"C": "=a25_oth", "D": "=a26_oth", **{c_(fy): "=oth_f" for fy in BU[1:]}}, fmt="0.00")
    row("etr", "有效稅率", "%", {"C": "=a25_etr", "D": "=a26_etr", **{c_(fy): "=tx_f" for fy in BU[1:]}}, fmt="0.0%")
    row("ni", "淨利（non-GAAP）", "US$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(fy): f"=({c_(fy)}{R['oi']}+{c_(fy)}{R['oth']})*(1-{c_(fy)}{R['etr']})" for fy in BU[1:]}})
    row("shb", "稀釋股數（不含客戶認股權證）", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(fy): f"={c_(fy - 1)}{r}*(1+sh_g)" for fy in BU[1:]}}, fmt="0.000")
    gwx = lambda fy: {2026: "gw_26", 2027: "gw_27", 2028: f"gw_28*um28_{s}", 2029: f"gw_29*um29_{s}"}[fy]
    row("wsh", "客戶認股權證稀釋股數（年中平均已歸屬 GW × 每 GW 股數）", "十億股", {"C": "=0", "D": "=0", **{c_(fy): f"=w_tot/w_gw*({gwx(fy - 1)}+{gwx(fy)})/2*{P['wf']}" for fy in BU[1:]}}, fmt="0.000", note="OpenAI、Meta 各 160M 股，依 GW 分批歸屬（Assumed 線性）；行使價 $0.01，視同全額稀釋")
    row("sh", "稀釋股數", "十億股", {c_(fy): f"={c_(fy)}{R['shb']}+{c_(fy)}{R['wsh']}" for fy in FYS}, fmt="0.000")
    row("eps", "EPS（non-GAAP）", "US$", {c_(fy): f"={c_(fy)}{R['ni']}/{c_(fy)}{R['sh']}" for fy in FYS}, fmt="0.00")
    row("epsg", "EPS 年增", "%", {c_(fy): f"={c_(fy)}{R['eps']}/{c_(fy - 1)}{R['eps']}-1" for fy in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — AMD 損益（US$B；non-GAAP；FY2025A 實績、FY2026 錨點，FY2027 起為模型）。會計年度＝日曆年。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026、FY2027 三情境相同（情境係數從 CY2028 才開始）"
PB = {}; r0 = 70
SCT = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCT[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("dc", "  Data Center", "US$B", "#,##0.0"), ("gpu", "    資料中心 GPU", "US$B", "#,##0.0"),
                                ("dcx", "    資料中心非 GPU", "US$B", "#,##0.0"), ("cg", "  Client＋Gaming", "US$B", "#,##0.0"), ("embr", "  Embedded", "US$B", "#,##0.0"),
                                ("gm", "毛利率", "%", "0.0%"), ("oi", "營業利益", "US$B", "#,##0.0"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利（non-GAAP）", "US$B", "#,##0.0"),
                                ("sh", "稀釋股數", "十億股", "0.000"), ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for fy in FYS:
            if key in ("g", "epsg") and fy == 2025: continue
            PL[f"{FC[fy]}{sr}"] = f"={FC[fy]}{PB[s][key]}"; PL[f"{FC[fy]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 由下而上營收、毛利率 − 錨點（k_rec、非 GPU 毛利率校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['bur']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：毛利率差"); PL[f"D{sr}"] = f"=D{PB['A']['bugm']}-a26_gp/a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
assert sr < 70, sr
PL.column_dimensions["A"].width = 60; PL.column_dimensions["B"].width = 14; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 13
eps = lambda s, y: f"PL!{FC[y]}{SUMR[(s, 'eps')]}"

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元；non-GAAP EPS）。目標價＝十年 non-GAAP P/E 中位數 × FY2027 EPS（各情境），情境加權（第 0 節）。結論另看「現價 ÷ FY2027、FY2028 EPS」。跨公司規則：合理倍數＝10 年（2016–2025 年底）P/E 的分位，依循環位置：見頂年 25 分位、中段中位數、谷底 75 分位；AMD 會計年度＝日曆年，用共用規則 classify（不平移），見頂年度讀設備市場模型"
V["A2"] = "依循環位置取分位的規則值：不適用（保留計算，不進結論；比照 NVIDIA 模型）。理由：歷史倍數是落後倍數（年底價 ÷ 當年 non-GAAP EPS），含高成長期溢價與 EPS 低谷年（2016 為虧損、2022 年 18.5x、2019 年 71.7x）；GAAP 倍數另受 Xilinx 併購攤銷扭曲（2023 年 283x），只列參考。本版未做先驗，無 Andy 合理 P/E 對照"
V.cell(4, 1, "舊法：規則值情境加權（參考）").font = BOLD
hdr(V, 5, ["項目", "值", "", "", "說明"])
V.cell(6, 1, "規則值情境加權價（參考，非目標價）"); V["B6"] = "=tp_wA*B8+tp_wB*B9+tp_wC*B10"; V["B6"].number_format = "0.00"
for c in range(1, 6): V.cell(6, c).fill = KEY
V.cell(7, 1, "相對現價（參考）"); V["B7"] = "=B6/px-1"; V["B7"].number_format = "0.0%"
for i, (s, nm) in enumerate((("A", "樂觀"), ("B", "基準"), ("C", "下檔")), 8):
    V.cell(i, 1, f"情境 {s} 隱含價（{nm}）"); V[f"B{i}"] = f"=pe_mid*{eps(s, 2027)}"; V[f"B{i}"].number_format = "0.00"
    V.cell(i, 5, f"十年 non-GAAP P/E 中位數 × 情境 {s} FY2027 EPS")
V.cell(11, 1, "情境權重 A／B／C"); V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0%"
V["E11"] = "In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"
V.cell(12, 1, "舊法說明"); V["B12"] = ("P/E：十年（2016–2025 年底）non-GAAP P/E 中位數（年底價 ÷ 公司 non-GAAP EPS，自算；規則值不適用時的替代）× FY2027（≈未來 12–14 個月）non-GAAP EPS，情境加權。"
                                       "偏誤方向：歷史倍數含 AMD 2017–2021 年轉型期與 AI 前的高成長溢價（偏高）。FY2027 EPS 三情境相同（情境係數自 CY2028 起），故三情境隱含價相同；"
                                       "樣板標籤 A 樂觀／B 基準／C 下檔沿用，本模型 A＝市場層基準、B＝AI 資本支出消化、C＝明顯下修")
V.cell(14, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 15, ["項目", "A", "B", "C", "說明"])
V.cell(16, 1, "現價（2026-10-23 收盤）"); V["B16"] = "=px"; V["B16"].number_format = "0.00"
for i, (lab, y) in enumerate((("現價 ÷ FY2026 EPS（錨點）", 2026), ("現價 ÷ FY2027 EPS（模型）", 2027), ("現價 ÷ FY2028 EPS（模型）", 2028), ("現價 ÷ FY2029 EPS（模型）", 2029)), 17):
    V.cell(i, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{i}"] = f"=px/{eps(s, y)}"; V[f"{c}{i}"].number_format = "0.0"
V.cell(21, 1, "現價 ÷ 共識 FY2027 EPS"); V["B21"] = "=px/cons_eps27"; V["C21"] = "=px/cons_eps27_lo"; V["B21"].number_format = "0.0"; V["C21"].number_format = "0.0"
V["E21"] = "B＝stockanalysis 15.76；C＝tickflow 四季合計 11.23（兩來源差大）"
V.cell(22, 1, "現價 ÷ 共識 FY2028 EPS（約 $20，二手）"); V["B22"] = "=px/cons_eps28"; V["B22"].number_format = "0.0"
V["E18"] = "FY2027 EPS 三情境相同（情境係數從 CY2028 才開始）"

V.cell(24, 1, "2. 循環位置（共用規則，不平移）與規則倍數").font = BOLD
V.cell(25, 1, "規則：各情境見頂年度（日曆年）取自設備市場模型（Mkt 頁「見頂年度」：" + "、".join(f"{s_} {PEAK[s_]}" for s_ in SC) + "）；AMD 會計年度＝日曆年：FY 早於見頂年＝中段（中位數）、等於＝見頂（25 分位）、晚於＝谷底（75 分位）。")
hdr(V, 26, ["情境", "設備市場見頂年度（CY）", "FY2027 判定", "FY2027 倍數", "FY2028 判定", "FY2028 倍數", "規則"])
for i, s in enumerate(SC, 27):
    V.cell(i, 1, s); V.cell(i, 2, f"=pk_{s}")
    V.cell(i, 3, "=" + CP.excel_rule("2027", f"pk_{s}"))
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, "=" + CP.excel_rule("2028", f"pk_{s}"))
    V.cell(i, 6, f'=IF(E{i}="見頂",pe_peak,IF(E{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 7, "tools/cycle_position.py classify（不平移）")
V.cell(31, 1, "3. 規則值隱含股價（美元）——不適用：落後倍數含高成長期溢價，保留計算但不進結論；規則值＝FY2027 EPS × 依循環位置的倍數").font = BOLD
hdr(V, 32, ["情境", "EPS 年度", "EPS", "倍數類型", "倍數", "隱含股價（不適用；約 FY2027 年底，未折現）", "相對現價", "折現到現在"])
vr = 33; VAL = {}
for j, s in enumerate(SC):
    for y, typ, nm in ((2027, f"D{27 + j}", "rule"), (2027, "pe_peak", "參考：25 分位"), (2027, "pe_mid", "參考：中位數"), (2027, "pe_trough", "參考：75 分位"), (2028, f"F{27 + j}", "rule28")):
        V.cell(vr, 1, s); V.cell(vr, 2, f"FY{y}"); V.cell(vr, 3, f"={eps(s, y)}").number_format = "0.00"
        V.cell(vr, 4, f'="規則值（不適用）："&C{27 + j}' if nm == "rule" else (f'="規則值（FY2028；不適用）："&E{27 + j}' if nm == "rule28" else nm))
        V.cell(vr, 5, f"={typ}").number_format = "0.0"
        V.cell(vr, 6, f"=C{vr}*E{vr}").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
        V.cell(vr, 8, f"=F{vr}/(1+r_disc)^(1.19+{y - 2027})").number_format = "#,##0"
        if nm in ("rule", "rule28"):
            for c in range(1, 9): V.cell(vr, c).fill = KEY
            VAL[(s, y)] = vr
        vr += 1
V.cell(vr, 1, "FY2028 折現年數＝1.19＋1；折現率為 In 頁 r_disc（Assumed 8–12%）。").font = Font(italic=True)
vr += 2
V.cell(vr, 1, "4. What must be true：要讓 FY2027 年底價等於現價 ×（1＋報酬），需要的倍數（情境 A 的 FY2027、FY2028 EPS）").font = BOLD; vr += 1
hdr(V, vr, ["FY2027 年底隱含價 ÷ 現價", "所需倍數（÷ FY2027 EPS）", "所需倍數（÷ FY2028 EPS）", "對照：歷史 25 分位", "對照：歷史中位數", "對照：歷史 75 分位"]); vr += 1
WMT0 = vr
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/{eps('A', 2027)}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/{eps('A', 2028)}").number_format = "0.0"
    V.cell(vr, 4, "=pe_peak").number_format = "0.0"; V.cell(vr, 5, "=pe_mid").number_format = "0.0"; V.cell(vr, 6, "=pe_trough").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝約兩年 10% 年報酬。所需倍數高於歷史中位數時，現價已預支 FY2028 以後的成長（大單放量、毛利率不被 HBM 壓縮）。")
WMT_NOTE = vr
vr += 2
assert vr < 70, vr
TBL = 112   # 自算 P/E 的輸入表（年底價、EPS）
PE_NG = [None] + [f"=ROUND(B{TBL + 2 + i}/C{TBL + 2 + i},2)" for i in range(1, 10)]
PE_GA = [None, 342.67, 57.69, 152.87, 44.96, 55.78, 73.60, 283.48, 122.01, f"=ROUND(B{TBL + 11}/D{TBL + 11},2)"]
block_end = VR.add_rule_block(V, 70, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "自算：年底收盤價 ÷ 公司 non-GAAP 稀釋 EPS", "vals": PE_NG, "tag": "Derived（價格：Macrotrends 與 digrin 一致；EPS：公司第四季新聞稿；2016 non-GAAP 虧損，不計）"},
    {"label": "Macrotrends GAAP P/E（年底價 ÷ TTM GAAP EPS；2025 自算）", "vals": PE_GA, "tag": "Interested-party（資料商）；GAAP 受 Xilinx 攤銷與 2020 稅務利益扭曲，只列參考，不採用", "use": False}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"={eps('A', 2027)}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"={eps('A', 2027)}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"={eps('A', 2027)}"}]}],
    "0.0", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位。本模型採用自算 non-GAAP P/E（9 年，2016 為虧損）；GAAP 序列（Macrotrends）只列參考。")
assert block_end < TBL, block_end
rule_rows = [rw for rw in range(70, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
V.cell(TBL, 1, "自算 P/E 的輸入：年底收盤價與公司 EPS（2016–2025）").font = BOLD
hdr(V, TBL + 1, ["年底", "收盤價（US$）", "non-GAAP 稀釋 EPS（公司）", "GAAP 稀釋 EPS（公司）", "來源"])
PXH = [11.34, 10.28, 18.46, 45.86, 91.71, 143.90, 64.77, 147.41, 120.79, 214.16]
NGH = [-0.14, 0.17, 0.46, 0.64, 1.29, 2.79, 3.50, 2.65, 3.31, 4.17]
GAH = [-0.60, 0.04, 0.32, 0.30, 2.06, 2.57, 0.84, 0.53, 1.00, 2.65]
for i in range(10):
    rw = TBL + 2 + i
    V.cell(rw, 1, str(2016 + i))
    for c, v in zip("BCD", (PXH[i], NGH[i], GAH[i])):
        V[f"{c}{rw}"] = v; V[f"{c}{rw}"].font = BLUE; V[f"{c}{rw}"].number_format = "0.00"
    V.cell(rw, 5, "價格：Macrotrends（2016–2024）與 digrin（2016–2025）；2025 僅 digrin 一個來源。EPS：AMD 第四季／全年新聞稿（sec.gov；2016–2017、2018–2019、2020–2021、2022–2023、2024–2025 稿）" if i == 0 else "同上")
V.column_dimensions["A"].width = 52
for c in "BCDEFGH": V.column_dimensions[c].width = 18
V.column_dimensions["E"].width = 40

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A）：對 FY2027、FY2028、FY2029 的營收、EPS 與隱含股價（十年 P/E 中位數 × EPS，與目標價同法）的影響。每個案例是一個完整的推導區塊（本頁下方，與 PL 同一引擎）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2027 隱含價（中位數 P/E）", "隱含價相對基準", "FY2028 營收", "FY2028 EPS", "FY2028 EPS 相對基準", "FY2028 隱含價（中位數 P/E）", "隱含價相對基準", "FY2029 EPS", "FY2029 EPS 年增（對 FY2028）"])
cases = [("AMD 顆數係數 低（×0.8）", "FY2027 起 AMD 認列顆數 ×0.8（大單延後）", dict(uf="uf_lo")),
         ("AMD 顆數係數 高（×1.2）", "×1.2（OpenAI、Meta、Anthropic、Oracle 提前放量）", dict(uf="uf_hi")),
         ("世代售價係數 高（隨 HBM 容量）", "FY2027 ×1.25、FY2028 起 ×1.5（Analogy）", dict(pg=lambda y: f"pg_{y % 100}_hi")),
         ("對照：FY2027 起每顆＝市場層 pa_amd（k_rec 只解釋 2026 售價）", "世代係數＝1 ÷ k_rec", dict(pg=lambda y: "1/k_rec")),
         ("HBM 吸收比例 0（全額轉嫁）", "吸收比例 0.0", dict(abs="abs_mem_lo", abs_s="abs_mem_lo")), ("HBM 吸收比例 1（全額吸收）", "吸收比例 1.0", dict(abs="abs_mem_hi", abs_s="abs_mem_hi")),
         ("GPU 毛利率 低（40%）", "FY2026 GPU 毛利率 40%（非 GPU 毛利率重新校準）", dict(gmg="gm_gpu_lo")), ("GPU 毛利率 高（55%）", "55%", dict(gmg="gm_gpu_hi")),
         ("資料中心 CPU 成長 低", "FY2027 +30%、FY2028 +10%、FY2029 +5%", dict(gcpu=lambda y: f"g_cpu{y % 100}_lo")),
         ("資料中心 CPU 成長 高（公司說法）", "FY2027 +70%、FY2028 +40%、FY2029 +25%", dict(gcpu=lambda y: f"g_cpu{y % 100}_hi")),
         ("營業費用率 −2pt", "營業費用額外 −2% × 營收", dict(opxa="opx_adj_lo")), ("營業費用率 +2pt", "+2% × 營收", dict(opxa="opx_adj_hi")),
         ("認股權證不歸屬", "OpenAI、Meta 認股權證 0 股", dict(wf="wf_lo")), ("認股權證歸屬加速（×1.5）", "已歸屬 GW ×1.5", dict(wf="wf_hi"))]
SE.cell(4, 1, "基準（情境 A）"); SE.cell(4, 2, "PL 區塊 A")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    SE[f"C{i}"] = f"={pre}E{R['rev']}"; SE[f"D{i}"] = f"={pre}E{R['eps']}"; SE[f"H{i}"] = f"={pre}F{R['rev']}"; SE[f"I{i}"] = f"={pre}F{R['eps']}"
    SE[f"M{i}"] = f"={pre}G{R['eps']}"; SE[f"N{i}"] = f"=M{i}/I{i}-1"; SE[f"M{i}"].number_format = "0.00"; SE[f"N{i}"].number_format = "0.0%"
    SE[f"F{i}"] = f"=D{i}*pe_mid"; SE[f"K{i}"] = f"=I{i}*pe_mid"
    SE[f"E{i}"] = f"=D{i}/D$4-1"; SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"J{i}"] = f"=I{i}/I$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "CH": SE[f"{c}{i}"].number_format = "#,##0.0"
    for c in "DI": SE[f"{c}{i}"].number_format = "0.00"
    for c in "FK": SE[f"{c}{i}"].number_format = "#,##0"
    for c in "EGJL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85；記憶體 B 路徑", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6；記憶體 C 路徑", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數、售價、吸收比例與 CPU 成長從 FY2027 起作用，情境 B、C 從 FY2028 起。隱含價欄＝EPS × 十年 non-GAAP P/E 中位數（與目標價同法），只反映 EPS 變動。")
SENS_NOTE = nn + 3
assert nn + 3 < 40
SE.column_dimensions["A"].width = 46; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKLMN": SE.column_dimensions[c].width = 15

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、大單、共識、市場層。只作事後對照，不為對上數字而調整參數。公司對 2027 的說法（資料中心逾一倍、伺服器 CPU 逾 70%）與共識只放在這裡，不用於校準"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4; RFMT = {}
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False):
    global rr_
    RFMT[rr_] = fmt
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
def sech(t):
    global rr_
    RC.cell(rr_, 1, t).font = BOLD; RFMT[rr_] = None; rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{SUMR[(s, key)]}")
eng = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
sech("A. FY2026 錨點與 FY2027–28 vs 公司說法與共識（校準時未使用）")
rec("FY2026 營收 vs 共識（S&P Global，stockanalysis）", "=cons_rev26", pl("rev", 2026), "stockanalysis（二手）", "另一來源 tickflow 47.15")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "stockanalysis（二手）", "tickflow 6.81", fmt="0.00")
rec("FY2027 資料中心營收 vs 公司「逾一倍」（FY2026 錨點 × 2）", "=a26_dc*(1+g_dc27)", pl("dc", 2027), "公司 Q2 2026 法說（二手）", "")
rec("FY2027 資料中心非 GPU 年增 vs 公司伺服器 CPU「逾 70%」（百分點）", "=g_srv27", lambda s: f"=PL!E{PB[s]['cpu']}/PL!D{PB[s]['cpu']}-1", "同上", "非 GPU 含 DPU、FPGA；基準 50%（Assumed）", fmt="0.0%", pt=True)
rec("FY2027 營收 vs 共識（stockanalysis「Revenue Next Year」）", "=cons_rev27", pl("rev", 2027), "stockanalysis（二手）", "tickflow 四季合計 67.1")
rec("FY2027 EPS vs 共識（stockanalysis 15.76）", "=cons_eps27", pl("eps", 2027), "stockanalysis（二手）", "差距超過 30% 的原因見 README", fmt="0.00")
rec("FY2027 EPS vs 共識（tickflow 四季合計 11.23）", "=cons_eps27_lo", pl("eps", 2027), "tickflow（二手；可能較舊）", "", fmt="0.00")
rec("FY2028 EPS vs 共識約 $20（Motley Fool，二手）", "=cons_eps28", pl("eps", 2028), "ts2.tech 轉述（二手）", "tickflow 四季合計 14.41", fmt="0.00")
rec("FY2029 EPS vs 公司策略期間 EPS 逾 $20", "=lt_eps", pl("eps", 2029), "公司 2025-11 分析師日（二手）", "公司 Q2 2026 稱將顯著超過", fmt="0.00")
rec("FY2027 資料中心 GPU 營收 vs 公司「2027 年資料中心 AI 數百億美元」（下限 20）", 20, pl("gpu", 2027), "公司分析師日（二手）", "")
rec("FY2027 毛利率 vs FY2026 錨點（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "Q 頁", "HBM 漲價（2027 +80%，三情境相同）× 吸收比例 0.5 壓 GPU 毛利率", fmt="0.0%", pt=True)
sech("B. 認列係數 k_rec 與每顆售價")
rec("k_rec（Derived）", None, lambda s: "=k_rec", "錨點 GPU ÷ 市場層顆數 × pa_amd", "FY2027 起沿用", fmt="0.000")
rec("FY2026 隱含每顆認列售價 vs HSBC MI355X（US$K）", "=p355", lambda s: "=p_real26", "HSBC（finviz 轉述，Interested-party）", "若 k_rec 全歸售價層；與 MI355X 價接近＝2026 的水準差主要是售價（MI355X 為主）", fmt="0.0")
rec("AMD 認列顆數 FY2026 vs 市場層（百萬顆）", "=u_amd_26", lambda s: f"=PL!D{PB[s]['urec']}", "markets/memory UNITS", "", fmt="0.00")
rec("FY2026 GPU 營收 vs FY2025（Assumed 6.5）年增", None, lambda s: f"=PL!D{PB[s]['gpu']}/PL!C{PB[s]['gpu']}-1", "Q 頁（皆含 Assumed 拆分）", "", fmt="0.0%")
sech("C. 大單（Interested-party）：認列顆數換算 GW")
rec("FY2027 AMD 認列顆數 ÷ 每 GW GPU 數 → 隱含 GW", "=gw_27-gw_26", lambda s: f"=PL!E{PB[s]['urec']}*1000000/gpw", "對照值＝In 頁假設的 2027 年新歸屬 GW", "每 GW 顆數以 VR200 NVL72 為 Analogy；每顆認列售價低於市場層時，顆數高估、GW 偏高", fmt="0.00")
rec("FY2028 隱含 GW", "=gw_28-gw_27", lambda s: f"=PL!F{PB[s]['urec']}*1000000/gpw", "同上", "", fmt="0.00")
rec("已宣布大單合計 GW（OpenAI 6＋Meta 6＋Anthropic 2＋Oracle 首批）", "=openai_gw+meta_gw+anth_gw+orcl_u*1000000/gpw", lambda s: f"=SUM(PL!E{PB[s]['urec']}:G{PB[s]['urec']})*1000000/gpw", "公司、Oracle 新聞稿", "模型值＝FY2027–29 認列顆數累計換算 GW；大單為多年期、上限值", fmt="0.0")
rec("認股權證稀釋占 FY2029 股數", None, lambda s: f"=PL!G{PB[s]['wsh']}/PL!G{PB[s]['sh']}", "In 頁 G 節", "上限 0.32B 股（約為 FY2026 股數的 19%）", fmt="0.0%")
sech("D. 份額（Units 頁，模組層級）")
rec("2027 年 AMD 金額份額", None, lambda s: f"=Units!D{SHR[s][('ds', 'amd')]}", "accel_share（NVIDIA 模型同一套價格）", "", fmt="0.0%")
rec("2027 年 AMD 顆數份額", None, lambda s: f"=Units!D{SHR[s][('us', 'amd')]}", "同上", "公司目標「雙位數份額」（2025 分析師日，二手）", fmt="0.0%")
rec("2027 年 NVIDIA 金額份額", None, lambda s: f"=Units!D{SHR[s][('ds', 'nv')]}", "同上", "與 NVIDIA 模型 Units 頁的 CY2027 值相同（同一份顆數與價格）", fmt="0.0%")
sech("E. 預留")
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 AMD 資料中心 GPU 營收 ÷ 部署 GW 對照 IT capex"); RC.cell(rr_, 10, "預留列：本版無數值"); RFMT[rr_] = None; rr_ += 1
RC_END = rr_
RC.column_dimensions["A"].width = 78
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 46; RC.column_dimensions["J"].width = 90

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝AMD 8-K 新聞稿、簡報與 10-Q，sec.gov 已逐頁抓取；法說逐字稿只經摘要）"; SO["A1"].font = BOLD
srcs = [("公司：Q2 2026 新聞稿（2026-08-04）", Q226), ("公司：Q2 2026 簡報（Q3 指引明細）", Q226S), ("公司：Q1 2026 新聞稿（2026-05-05）", Q126),
        ("公司：Q4 2025 與全年新聞稿（2026-02-03）", Q425), ("公司：Q2 2026 10-Q（股數、買回）", E8K + "000000248826000123/amd-20260627.htm"),
        ("公司：OpenAI 6 GW 合作新聞稿（2025-10-06）", "https://ir.amd.com/news-events/press-releases/detail/1260/amd-and-openai-announce-strategic-partnership-to-deploy-6-gigawatts-of-amd-gpus"),
        ("公司：Meta 6 GW 合作新聞稿（2026-02-24，8-K EX-99.1）", E8K + "000000248826000045/pressreleasedatedfebruary2.htm"),
        ("Oracle：5 萬顆 MI450 新聞稿（2025-10-14）", "https://oracle.com/news/announcement/ai-world-oracle-and-amd-expand-partnership-to-help-customers-achieve-next-generation-ai-scale-2025-10-14"),
        ("公司：歷年第四季新聞稿（EPS 歷史）", E8K + "000000248824000009/q42023991final.htm；000000248822000014/q421991amd.htm；000000248818000014/q417991.htm；000000248820000006/amdq4andfy2019earningsslides.pdf"),
        ("公司：Q3 2026 財報日（2026-11-03 盤後）", "https://finviz.com/news/399552/amd-to-report-fiscal-third-quarter-2026-financial-results"),
        ("二手：MarketBeat Q2 2026 法說摘要", "https://www.marketbeat.com/instant-alerts/advanced-micro-devices-q2-earnings-call-highlights-2026-08-04/"),
        ("二手：platformaeronaut Q2 2026 摘要", "https://transcripts.platformaeronaut.com/summaries/AMD-2Q26-AI-Summary"),
        ("二手：Motley Fool Q4 2025 逐字稿（分析師日目標重述）", "https://www.fool.com/earnings/call-transcripts/2026/02/03/amd-amd-q4-2025-earnings-call-transcript/"),
        ("二手：Next Platform（Q4 2025 Instinct／EPYC 拆分估計）", "https://www.nextplatform.com/2026/02/04/amd-finally-makes-more-money-on-gpus-than-cpus-in-a-quarter"),
        ("二手：stockanalysis（現價、共識）", "https://stockanalysis.com/stocks/amd/forecast/"), ("二手：tickflow（共識，各季）", "https://www.tickflow.io/stock/AMD/forecast"),
        ("二手：ts2.tech（2028 EPS 約 $20，轉述 Motley Fool）", "https://ts2.tech/en/amd-nasdaqamd-shares-require-43-times-earnings-multiple-for-price-to-double-by-2028/"),
        ("二手：Macrotrends（年底價、GAAP P/E）", "https://www.macrotrends.net/stocks/charts/AMD/amd/pe-ratio"), ("二手：digrin（年底價）", "https://www.digrin.com/stocks/detail/AMD/price"),
        ("二手：heygotrade（每日收盤報導）", "https://www.heygotrade.com/en/news/amd-amd-stock-news-october-20-2026/"),
        ("市場層：markets/memory/market.py（UNITS）；tsmc、nvidia（accel_share、In 頁價格與 HBM 路徑）、markets/chip_bridge、markets/equipment 成品（本 repo）", ""),
        ("未找到：GPU／CPU 拆分（公司不揭露）；認股權證會計處理與各批股數；MI450 每顆售價；Q2 2026 法說逐字稿原文", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 80; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("AMD（AMD）收支與估值模型 v0.2.2（AI 半導體 Project：運算層，GPU 側非 NVIDIA 份額）　2026-10-25", True),
 ("本版：運算層第三個模型，複製 nvidia/ 的結構，匯入 nvidia/accel_share.py 的份額函式（AMD 為 GPU 側非 NVIDIA 份額）；Val 頁第 0 節為情境加權目標價。", False),
 ("模型命題：AMD 的資料中心 GPU 營收＝加速器市場（本專案加速器數量層 × 各平台每顆價）× AMD 金額份額；OpenAI、Meta 各 6 GW 與 Anthropic、Oracle 大單在 2026 下半年起放量，但每 GW 附帶認股權證稀釋，HBM 漲價又壓 GPU 毛利率。股價約為 FY2027 共識 EPS 的 33 倍：要多少大單如期放量、毛利率守住多少，才撐得住現價。", False),
 ("核心問題：以市場層顆數推得的 AMD 資料中心 GPU 營收與 EPS，能否支撐現價與公司「2027 年資料中心逾一倍、EPS 逾 $20」的說法；差距來自顆數、售價還是毛利率。", False), ("", False),
 ("會計年度：AMD 會計年度約在 12 月最後一個週六結束，等於日曆年（FY2026＝CY2026），上游市場層直接對照，不平移。預測期 FY2027–FY2029；FY2026 為錨點（Q1、Q2 實績＋Q3 指引＋Q4 假設）。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("資料中心 GPU：加速器市場金額（Units 頁，markets/memory UNITS × 情境係數 × NVIDIA 模型同一套每顆價）× AMD 金額份額 × k_rec（FY2026 校準）× 世代售價係數（基準 1.0）；HBM 漲價依吸收比例轉嫁或壓毛利率。資料中心非 GPU（EPYC CPU 等）、Client、Gaming、Embedded：FY2026 錨點 × 簡化成長率（Assumed）。", False),
 ("毛利率：GPU 毛利率（Assumed 50%）＋非 GPU 毛利率（FY2026 校準到錨點，Derived）；FY2027 起變化來自 GPU 占比上升與 HBM 價格（markets/memory 2027 +80%）。營業費用彈性（FY2026 推得）、稅率 13%（公司）、股數＋OpenAI／Meta 認股權證（各 160M 股，依 GW 線性歸屬，Assumed）。", False),
 ("EPS 口徑：non-GAAP（公司口徑，排除股票薪酬與併購攤銷）；共識同口徑。", False),
 ("校準：k_rec＝FY2026 錨點 GPU 營收 ÷（AMD 市場層 CY2026 顆數 × pa_amd）；錨點 GPU 本身含 Assumed 拆分（公司不揭露）。不用共識、不用公司對 2027 的說法；後者只在 Recon 對帳。", False),
 ("估值：目標價＝十年 non-GAAP P/E 中位數 × FY2027 EPS（各情境），權重 A 0.3／B 0.5／C 0.2（Andy 可改）；依循環位置取分位的規則值不適用（落後倍數含高成長期溢價，比照 NVIDIA），保留計算。循環位置（共用規則，不平移）：見頂年度 " + "、".join(f"{s_} {PEAK[s_]}" for s_ in SC) + " → FY2027：" + "、".join(f"{s_} {CP.classify(2027, PEAK[s_])}" for s_ in SC) + "。本版未做先驗，無 Andy 合理 P/E 對照。", False), ("", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文＝8-K 新聞稿、簡報與 10-Q（sec.gov，已抓取）；法說逐字稿只經摘要（標二手）。", False),
 ("已知限制：(1) 公司不拆分 GPU 與 CPU，FY2026 錨點的 GPU 營收（約 15）為 Assumed 拆分，k_rec 隨之變動。(2) k_rec 是顆數層與售價層的合併係數；FY2027 起沿用（世代售價係數 1.0）可能低估 MI450 的售價提升，Sens 列「隨 HBM 容量」與「每顆＝市場層價」兩個上界。(3) HBM 吸收比例、GPU 毛利率、各業務成長率為 Assumed。(4) 認股權證的會計處理（是否抵減營收）與各批股數未找到；本模型只計股數稀釋。(5) 情境 B、C 只影響加速器顆數與 HBM 價格，CPU 等業務三情境相同。(6) 共識兩來源差距大（FY2027 EPS 15.76 vs 11.23）。(7) 十年 P/E 為落後倍數，含高成長期溢價（目標價方法已註明偏誤方向）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ---------------- v0.2：三情境延伸到 FY2032、Val 第 0 節（DCF 與 Comparable 加權）、Sens 目標價參數 ----------------
XH, XR, XC = TP.extend_pl(PL, PB, FC, FYS, 2029, {"rev": "rev", "ni": "ni", "sh": "sh", "eps": "eps"}, "US$",
    "三情境延伸到 FY2032（v0.2；FY2025A–FY2029 引用上方推導區塊，FY2030 起為收斂法，非驅動因子推導；收斂年度為橘底）",
    "上游市場層（markets/*、台積電顆數）只到 CY2029＝FY2029；FY2030 起：營收年增由 FY2029 年增線性收斂到 FY2032 的長期成長（A＝tp_g＋tp_gd、B＝tp_g、C＝tp_g−tp_gd），"
    "淨利率由 FY2029 線性收斂到中段利潤率（A＝FY2029、B＝FY2029 與 C 目標的中點、C＝FY2029 × nm10_k；同口徑十年平均未找到），股數沿用 FY2029 年變動")
# v0.2.1：收斂目標淨利率改為 A ≥ B ≥ C（v0.2 的 A＝FY2029 淨利率，因 HBM 成本 A 最高而低於 B 目標）。
#   C＝FY2029(C) × nm10_k；B＝MAX(FY2029(B) 與 C 規則的中點, C 目標)；A＝MAX(FY2029(A), B 目標)
_lc = XC[2029]; _nm = {s_: f"${_lc}${XR[(s_, 'nm')]}" for s_ in SC}
_T = {"C": f"{_nm['C']}*nm10_k"}
_T["B"] = f"MAX(({_nm['B']}+{_nm['B']}*nm10_k)/2,{_T['C']})"
_T["A"] = f"MAX({_nm['A']},{_T['B']})"
_conv = [y for y in range(2030, 2033)]
for s_ in ("A", "B"):
    for k_, y in enumerate(_conv, 1):
        PL[f"{XC[y]}{XR[(s_, 'nm')]}"] = f"={_nm[s_]}+({_T[s_]}-{_nm[s_]})*{k_}/{len(_conv)}"
    PL.cell(XR[(s_, "nm")], 3 + len(XC), {"A": "FY2030 起：FY2029 淨利率線性收斂到 MAX(FY2029 淨利率, B 目標)（v0.2.1：確保收斂目標 A ≥ B ≥ C）",
                                         "B": "FY2030 起：FY2029 淨利率線性收斂到 MAX(FY2029 與 FY2029 × nm10_k 的中點, C 目標)（v0.2.1：確保 B ≥ C）"}[s_])
_h = next(r for r in range(1, V.max_row + 1) if V.cell(r, 1).value == "P/E：十年歷史倍數")
TPL, TPN = TP.add_tp_section(wb, V,
    eps={s: [f"PL!{XC[y]}{XR[(s, 'eps')]}" for y in range(2027, 2033)] for s in SC}, e_mon=12,
    fwd={s: [(f"PL!{XC[2027]}{XR[(s, 'eps')]}", "2/12"), (f"PL!{XC[2028]}{XR[(s, 'eps')]}", "10/12")] for s in SC},
    fwd_note="會計年度＝日曆年：約 12 個月後（2027-10 底）的未來 12 個月＝2027-11～2028-10（2/12 × FY2027＋10/12 × FY2028）",
    peers=[("NVIDIA", "NVDA", 17.44, "2026-10-23", "stockanalysis.com 統計頁（Forward PE）", "Interested-party（資料商彙整賣方預估）"),
           ("Broadcom", "AVGO", 18.52, "2026-10-23", "同上", "Interested-party（同上）"),
           ("Intel", "INTC", 50.64, "2026-10-23", "同上", "Interested-party（同上）"),
           ("Qualcomm", "QCOM", 18.98, "2026-10-23", "同上", "Interested-party（同上）"),
           ("Marvell", "MRVL", 40.44, "2026-10-23", "同上", "Interested-party（同上）")],
    peer_note="券商常用的 CPU／GPU／AI 半導體同業（NVIDIA、Intel 為直接競爭者）；不含 A 股（Andy 11:53）。前瞻 P/E 口徑為資料商的未來 12 個月共識 EPS，各家 GAAP／non-GAAP 口徑不一",
    cons_tp="cons_tp", cons_note="stockanalysis.com：55 位分析師平均目標價（2026-10-25 頁面）；Interested-party",
    cons_fwd="2/12*cons_eps27+10/12*cons_eps28", cons_fwd_note="同權重的共識 EPS：2/12 × 共識 FY2027（15.76，stockanalysis）＋10/12 × 共識 FY2028（約 20，二手）",
    eps28={s: f"PL!{XC[2028]}{XR[(s, 'eps')]}" for s in SC}, hist=f"B{_h + 2}:B{_h + 11}", fmt="#,##0.00",
    method="DCF：股東現金流＝non-GAAP EPS ×（1 − tp_g ÷ tp_roe），tp_r 10%、tp_g 4%、tp_roe 0.20（工作單預設；GAAP ROE 受 Xilinx 攤銷壓低不採）；基準日 2026 年底，FY2027 起逐年折現，FY2032 後永續。"
           "Comparable：NVIDIA、Broadcom、Intel、Qualcomm、Marvell 前瞻 P/E 中位數 × 12 個月後前瞻 EPS。偏誤方向：non-GAAP EPS 排除股票薪酬與攤銷（偏高）；C 情境 FY2030 起淨利率收斂到 FY2029 × 0.8（Assumed）；"
           "歷史分位用自算十年 non-GAAP 落後 P/E（含轉型期高倍數）；永續價值占 DCF 價值過半，對 tp_r、tp_g 敏感（Sens 頁）。")
assert nn + 9 < 40, nn
TP.add_tp_sens(SE, nn + 5, TPL, "#,##0.00")
TP.dump_rowmap(TPL, TPN, XH, XR, XC, nn + 5)
TP.yellow_inputs(IN)
RD.cell(len(lines) + 1, 1, "v0.2（2026-10-25）：(1) 目標價改為 DCF 與 Comparable 加權（Val 頁最前面「0. 目標價」：0.1 DCF 股東權益現金流、0.2 Comparable 同業前瞻 P/E、0.3 結論；Val 既有內容整體下移，公式一併改寫）；"
        "v0.1 的「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留。(2) PL 頁末新增三情境延伸到 FY2032：FY2030 起為收斂法，非驅動因子推導（上游市場層只到 CY2029＝FY2029）。"
        "(3) 現價維持 2026-10-23 收盤 516.79（補第二來源 digrin）。(4) In 頁 J 節新增目標價參數（黃底，Andy 可改）；Sens 頁新增 tp_r、tp_g、tp_w_dcf 對加權目標價的影響。FY2025–FY2029 既有數字不變。")
RD.cell(len(lines) + 2, 1, "情境對應（v0.2 延伸期）：A＝順利（市場層基準顆數）、B＝平平（AI 資本支出消化：CY2028 ×0.8、CY2029 ×0.85）、C＝不如預期（CY2028–29 ×0.6）；FY2027 三情境相同，情境差從 FY2028 起。本版未做先驗，無 Andy 合理 P/E 對照。")

RD.cell(len(lines) + 3, 1, "v0.2.1（2026-10-25）情境一致性修正：(1) HBM 漲價吸收比例拆成 abs_mem_A／_B／_C（In 頁 K 節：A 0.2、B 0.5、C 0.7，A 轉嫁最多），FY2028 起讀本情境、FY2026–27 三情境共用 abs_mem 0.5（FY2027 三情境仍相同）；"
        "A 先試 0.3 時 FY2028 EPS 仍 B＞A（認股權證稀釋 A 最多），故取 0.2。認股權證照原邏輯。(2) FY2030–32 收斂目標淨利率改為 A ≥ B ≥ C：C＝FY2029(C) × nm10_k、B＝MAX(B 中點, C 目標)、A＝MAX(FY2029(A), B 目標)。修正後 FY2028–FY2032 EPS A ≥ B ≥ C。")

# ---------------- v0.2.2：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00、14:17）；各頁只在現有最後一列之後追加，不位移既有列 ----------------
_up = next(r for r in range(1, V.max_row + 1) if V.cell(r, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 2; REC_ROW = GRADE_ROW + 1
V.cell(GRADE_ROW, 1, "上檔空間級距")
V.cell(GRADE_ROW, 2, f'=IF(B{_up}>0.5,"大於 50%",IF(B{_up}>0.25,"25%–50%",IF(B{_up}>=0,"0%–25%",IF(B{_up}>=-0.25,"0%～−25%",IF(B{_up}>=-0.5,"−25%～−50%","小於 −50%")))))')
V.cell(GRADE_ROW, 5, f"以「上檔空間：加權目標價」（B{_up}）分六級（Andy 2026-10-25 14:00）")
V.cell(REC_ROW, 1, "投資建議")
V.cell(REC_ROW, 2, f'=IF(B{_up}>0.5,"強力買進",IF(B{_up}>0.25,"買進",IF(B{_up}>=-0.25,"中立",IF(B{_up}>=-0.5,"放空","強力放空"))))')
V.cell(REC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = ['目標價 US$177（DCF 152、Comparable 202 各半），比現價 516.79（2026-10-23）低 65.7%，級距 小於 −50%；Comparable 單獨為 −60.9%。EPS 與倍數兩個因素都在：本模型 12 個月後前瞻 EPS 10.65 比共識前瞻 EPS 19.29 低 44.8%；同業（NVIDIA、Broadcom、Intel、Qualcomm、Marvell）前瞻 P/E 中位數 19.0 倍，現價隱含 26.8 倍共識 EPS（48.5 倍本模型 EPS），共識目標價 637.61 隱含 33.0 倍（比同業高 74.1%）。', 'EPS 缺口從 FY2027 就出現：本模型 FY2027 EPS 9.14（三情境相同）比共識 15.76 低 42.0%，FY2028 B 情境 11.11 約為共識（約 20，二手）的 56%；單一變數補不了缺口——Sens 頁 AMD 顆數 ×1.2 只讓 FY2027 EPS +5.8%、世代售價隨 HBM 容量上調 +7.2%，賣方路徑隱含大客戶訂單全面放量。', 'HBM 漲價吸收比例是最大單一變數（全額轉嫁時 FY2027 EPS +10.9%、全額吸收時 −10.9%），其次是營業費用率 ±2pt（∓8.5%）（Sens 頁）。下一個檢驗點是 Q3 2026 財報（2026-11-03 盤後）的資料中心營收與 MI450 出貨時程。']   # 重點觀察（與 README.md 相同文字；數字取自重算後的 Val／PL／Sens 頁）
_rr = RD.max_row + 1
RD.cell(_rr, 1, 'v0.2.2（2026-10-25）：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁末新增「上檔空間級距」「投資建議」兩列（Andy 2026-10-25 14:00、14:17 固定規則），README 頁末新增投資建議與重點觀察三點；既有儲存格位置與數值不變。')
RD.cell(_rr + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"').font = Font(bold=True)
for _j, _t in enumerate(OBS, 1): RD.cell(_rr + 1 + _j, 1, f"重點觀察 {_j}：" + _t)
if os.environ.get("RATING_ROWMAP"):   # 除錯用：輸出新增列號（供撰寫 html.json）；CI 不設此變數
    print("ROWMAP", GRADE_ROW, REC_ROW, _rr, _rr + 1, _rr + 2)


order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# 相對倍數調整（Andy 2026-10-27「照建議做」；tools/relative_multiple.py，2026-10-29 接入）
import os, sys  # noqa: E401
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools"))
import relative_multiple as RMX  # noqa: E402
RMX.apply(wb, "amd")
wb.save(sys.argv[1])
if os.environ.get("AMD_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "VAL": {f"{k[0]}|{k[1]}": v for k, v in VAL.items()}, "CHK": CHK_ROW, "CAL": CAL,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RFMT": {str(k): v for k, v in RFMT.items()}, "RC_END": RC_END, "INROW": INROW, "WMT0": WMT0, "WMT_NOTE": WMT_NOTE,
               "VR_END": vr, "RULE": rule_rows}, open(os.environ["AMD_ROWMAP"], "w"))
