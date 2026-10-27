# Dell Technologies（DELL US）收支與估值模型 v0.1（AI 半導體 Project：AI 伺服器品牌次產業，樣板＝wiwynn/；smci/ 尚未在 main）
# 用法：python3 dell/build_dell.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度：FY(t)＝t−1 年 2 月～t 年 1 月底（FY2027：2026-01-31～2027-01-29），11/12 落在日曆年 t−1，所以 FY(t) 直接取 CY(t−1)
#   （與 NVIDIA、Marvell 模型相同；不跨年加權）。上游顆數只到 CY2029＝FY2030；FY2031 起為收斂法。
# 金額：十億美元（US$B）；EPS：美元。口徑：non-GAAP（公司指引、共識與同業前瞻 P/E 皆為 non-GAAP；排除攤銷、股票薪酬、重組與投資損益）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
QU_F = only("quanta/dist/*.xlsx"); WI_F = only("wistron/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith("um") or str(k).startswith("u_") or str(k).startswith("g29_")):
        TS[k] = (r[4].value, r[0].row)
GCL = ("bw", "hp"); VCL = ("rb", "ru", "amd")          # 商用 GPU：Blackwell／Hopper 用 GB300 每顆價值；Rubin／Rubin Ultra 用 VR200；AMD 2026（MI355）計 GB300、2027 起（MI450 整櫃）計 VR200（Siliconomics Rack 頁，CF10）
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
assert all(f"u_{p}_{y}" in TS for p in GCL + VCL for y in (26, 27, 28)) and all(f"g29_{p}" in TS for p in GCL + VCL), TS.keys()
def units(plats):
    u = {y: sum(TS[f"u_{p}_{y}"][0] for p in plats) for y in (26, 27, 28)}
    u[29] = sum(TS[f"u_{p}_28"][0] * (1 + TS[f"g29_{p}"][0]) for p in plats)
    return u
UG = units(GCL); UV = units(("rb", "ru")); UA = units(("amd",))
UG[26] += UA[26]
for y in (27, 28, 29): UV[y] += UA[y]
# CF10：商用 GPU 系統市場只有一個值（Siliconomics Rack 頁）；本模型的世代分組必須重現同一市場（情境 A，未四捨五入）
def _chk_mkt():
    for j, y in enumerate((26, 27, 28, 29)):
        _m = UG[y] * BR["rack"][0][1] / BR["gpu_n"][0][1] * 1000 + UV[y] * BR["rack"][0][4] / BR["gpu_n"][0][4] * 1000
        assert abs(_m - SIL.NAMES["SIF_GPUSysMkt_A"][j]) < 1e-6, ("商用 GPU 系統市場與 Siliconomics 不一致", y, _m, SIL.NAMES["SIF_GPUSysMkt_A"][j])
U_SRC = f"{rel(TSMC_F)}｜In!E{TS['u_bw_26'][1]}:E{TS['g29_amd'][1]}"
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("rack", "gpu_n"):
        BR[k] = ([r[j].value for j in range(3, 9)], r[0].row)
assert len(BR) == 2, BR.keys()
_chk_mkt()
BR_SRC = f"{rel(BR_F)}｜In!D{BR['rack'][1]}:I{BR['rack'][1]}（機架單價）÷ In!E{BR['gpu_n'][1]}（每架 GPU 數）"
def pu(j):   # 每顆系統價值（US$K/顆）＝機架單價（US$M）÷ 每架 GPU 數 × 1000；j：0–2 GB300 低基高、3–5 VR200 低基高
    return round(BR["rack"][0][j] / BR["gpu_n"][0][j] * 1000, 6)
# 組裝廠模型（只作份額合理性對照，Recon）
QMK = None
for r in openpyxl.load_workbook(QU_F, data_only=True)["Units"].iter_rows(min_row=4):
    if r[0].value == "AI 伺服器組裝市場（ODM 價）":
        QMK = ([r[2 + j].value for j in range(4)], f"{rel(QU_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert QMK and all(v is not None for v in QMK[0]), QMK
WSUM = None
for r in openpyxl.load_workbook(WI_F, data_only=True)["Units"].iter_rows(min_row=4):
    if r[0].value and str(r[0].value).startswith("鴻海＋廣達＋緯創（含緯穎）合計有效份額"):
        WSUM = ([r[2 + j].value for j in range(4)], f"{rel(WI_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert WSUM and all(v is not None for v in WSUM[0]), WSUM

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；每顆系統價值：千美元／顆；比例：%。口徑：non-GAAP。每個輸入盡量附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司原文＝Dell 8-K 財報新聞稿（SEC）；公司指引與展望為 Interested-party；券商與資料商估計皆為 Interested-party。加速器顆數讀 tsmc 成品、每顆系統價值讀橋接表（唯讀）"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
K4 = "https://www.sec.gov/Archives/edgar/data/1571996/000157199626000003/exhibit991earnings8kq4fy26.htm"
K1 = "https://www.sec.gov/Archives/edgar/data/0001571996/000157199626000021/exhibit991earnings8kq1fy27.htm"
K2 = "https://www.sec.gov/Archives/edgar/data/0001571996/000157199626000039/exhibit991earnings8kq2fy27.htm"
PUL = "https://pulse2.com/dell-technologies-q2-revenue-hits-record-47-billion-as-ai-server-backlog-reaches-95-billion-and-fy27-outlook-rises-to-192-billion/"
SNL = "https://www.storagenewsletter.com/2026/09/03/dell-technologies-fiscal-2q27-financial-results/"
ZK = "https://www.zacks.com/stock/news/2984030/dell-q2-earnings-beat-estimates-strong-ai-demand-aids-revenue-growth"
SA = "https://stockanalysis.com/stocks/dell/statistics/"
SAF = "https://stockanalysis.com/stocks/dell/forecast/"
MS = "https://uk.marketscreener.com/quote/stock/DELL-TECHNOLOGIES-INC-50061235/consensus/"
ALP = "https://alphaspread.com/security/nyse/dell/analyst-estimates"
INV = "https://www.investing.com/equities/dell-inc-consensus-estimates"
CMC = "https://companiesmarketcap.com/dell/pe-ratio/"
MT = "https://www.macrotrends.net/stocks/charts/DELL/dell/pe-ratio"
SAP = "https://stockanalysis.com/stocks/{}/statistics/"
URLNAME = {K4: "Dell 8-K 4Q FY26（2026-02-26）", K1: "Dell 8-K 1Q FY27（2026-05）", K2: "Dell 8-K 2Q FY27（2026-09-01）", PUL: "Pulse 2.0 2026-09-01",
           SNL: "StorageNewsletter 2026-09-03", ZK: "Zacks 2026-09", SA: "stockanalysis.com（統計）", SAF: "stockanalysis.com（預估）", MS: "MarketScreener",
           ALP: "Alpha Spread", INV: "Investing.com", CMC: "companiesmarketcap", MT: "Macrotrends"}

sec("A. 加速器顆數與每顆系統價值（讀 tsmc、markets/chip_bridge 成品；情境 A 基準，B／C 依同字母顆數係數縮放）")
for y in (26, 27, 28, 29):
    inp(f"ug_{y}", f"GB300 級商用 GPU 顆數 CY20{y}（Blackwell＋Hopper" + ("＋AMD MI355）" if y == 26 else "）"), "百萬顆", None, round(UG[y], 6), None, "Derived", U_SRC if y == 26 else "同上",
        "tsmc 成品 u_bw、u_hp（2026 另含 u_amd）；CY2029＝CY2028 × (1＋各類成長)；世代分組與 Siliconomics Rack 頁、美超微模型相同（CF10）" if y == 26 else "同上")
for y in (26, 27, 28, 29):
    inp(f"uv_{y}", f"VR200 級商用 GPU 顆數 CY20{y}（Rubin＋Rubin Ultra" + ("）" if y == 26 else "＋AMD MI450）"), "百萬顆", None, round(UV[y], 6), None, "Derived", U_SRC if y == 26 else "同上",
        "tsmc 成品 u_rb、u_ru、u_amd；Google TPU、Trainium 與其他 ASIC 不含（雲端自建，Dell 不是主要供應商）" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：CY2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", f"{rel(TSMC_F)}｜In!E{TS[f'um28_{s}'][1]}", "A＝順利（基準）；B＝平平；C＝不如預期；與台積電、廣達、緯穎模型同字母")
    inp(f"um29_{s}", f"情境 {s}：CY2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", f"{rel(TSMC_F)}｜In!E{TS[f'um29_{s}'][1]}", "同上")
inp("v_gb", "每顆系統價值（GB300 NVL72 機架價 ÷ 72）", "US$K/顆", pu(0), pu(1), pu(2), "Derived", BR_SRC, "Tokenomics v20.1 機架單價（經橋接表；唯讀）；Dell 為 OEM，營收含整櫃（含 GPU）")
inp("v_vr", "每顆系統價值（VR200 NVL72 機架價 ÷ 72；AMD Helios 假設相同）", "US$K/顆", pu(3), pu(4), pu(5), "Derived", BR_SRC, "同上；AMD MI450／Rubin Ultra 每顆價值未另訂（Assumed：同 VR200）")
inp("mf", "AI 部署係數（Sens 用；FY2028 起）", "x", 0.8, 1.0, 1.2, "Assumed", "無", "±20%")

sec("B. Dell 份額（FY2027 以公司 AI 伺服器營收指引校準）")
inp("rel_28", "Dell 份額相對值 FY2028 ÷ FY2027（全部情境）", "x", 0.75, 0.90, 1.05, "Assumed", K2 + "：2Q FY27 AI 訂單 609 億美元、在手訂單 950 億美元（公司）", PUL + "：近四季 AI 訂單 1,317 億美元；份額下檔：雲端與 neocloud 直接向 ODM 採購、Supermicro／HPE 競價")
inp("rel_29", "Dell 份額相對值 FY2029 ÷ FY2027", "x", 0.65, 0.85, 1.0, "Assumed", "同上；客戶集中於 neocloud 與大型 AI 實驗室（xAI、CoreWeave 等，媒體）", "—")
inp("rel_30", "Dell 份額相對值 FY2030 ÷ FY2027", "x", 0.55, 0.80, 1.0, "Assumed", "同上", "—")
inp("srel_A", "情境 A：FY2029 起份額額外乘數", "x", None, 1.0, None, "Assumed", "公司特有變數依情境調整，方向 A 最有利", "—")
inp("srel_B", "情境 B：FY2029 起份額額外乘數", "x", None, 0.95, None, "Assumed", "AI 資本支出消化時，neocloud 客戶（信用較弱）先縮", "—")
inp("srel_C", "情境 C：FY2029 起份額額外乘數", "x", None, 0.85, None, "Assumed", "同上，下修更深", "—")

sec("C. FY2027 錨點（1H FY27 實績＋公司全年指引＋2H 假設）")
inp("g_rev", "FY2027 營收指引", "US$B", 190.0, 192.0, 194.0, "Interested-party", K2 + "：FY27 營收 1,920 億美元（由 1,670 上修）", ZK + "：±20 億美元")
inp("g_ai", "FY2027 AI 伺服器營收指引", "US$B", None, 74.0, None, "Interested-party", K2 + "：FY27 AI-Optimized Servers 約 740 億美元（由 600 上修）", PUL + "：同")
inp("g_eps", "FY2027 non-GAAP EPS 指引", "US$", 25.25, 25.50, 25.75, "Interested-party", K2 + "：non-GAAP 稀釋 EPS 25.50（GAAP 24.37）", ZK + "：±0.25")
inp("g_q3rev", "3Q FY27 營收指引（參考）", "US$B", 48.5, 49.0, 49.5, "Interested-party", K2 + "：3Q FY27 營收 490 億美元、non-GAAP EPS 6.50", ZK + "：±5 億美元")
inp("st_h2", "2H FY27 儲存營收 ÷ 1H", "x", 1.05, 1.12, 1.20, "Assumed", SNL + "：3Q 儲存年增約 mid-teens、全年約 200 億美元（媒體轉述）", "1H FY27 91.84 億 × 1.12 ＝ 2H 102.9 億（全年 194.7 億）")
inp("csg_h2", "2H FY27 CSG 營收 ÷ 1H − 1", "%", -0.03, 0.03, 0.08, "Assumed", "CSG 2Q FY27 年增 20%、季增 6%（公司）", "公司未給 CSG 全年指引：未找到")
inp("isg_h2m", "2H FY27 ISG 營業利益率", "%", 0.12, 0.135, 0.155, "Assumed", K2 + "：2Q FY27 ISG 營益率 15.0%、1H 12.9%（公司）", "AI 比重上升稀釋 ISG 營益率（Assumed）")
inp("csg_h2m", "2H FY27 CSG 營業利益率", "%", 0.065, 0.075, 0.08, "Assumed", K2 + "：2Q FY27 CSG 營益率 7.6%、1H 7.8%（公司）", "FY26 5.6%")
inp("h2_sh", "2H FY27 稀釋股數（3Q、4Q 平均）", "十億股", 0.635, 0.642, 0.650, "Assumed", K2 + "：2Q 稀釋股數 6.52 億、1H 6.54 億；2Q 買回 950 萬股（均價 401 美元）", SA + "：流通股數 6.3581 億股（2026-10-23）")
inp("m_nai", "非 AI ISG（傳統伺服器、網通、儲存）營業利益率（拆分 AI 毛利用）", "%", 0.14, 0.16, 0.19, "Assumed", "公司不揭露 AI 伺服器營益率：未找到", "FY26：ISG 營益率 11.7%；非 AI 取 16% 時 AI 約 5%（Q 頁檢查列）")

sec("D. 營業利益率（FY2028 起）")
inp("ai_dm", "AI 伺服器營業利益率變動（FY2028 起，相對 FY2027 校準值；全部情境）", "pt", -0.03, -0.01, 0.01, "Assumed", "Rubin 機架單價較高、GPU 轉手比重高，營益率（%）下降；規模效益為上檔", "—")
inp("dm_A", "情境 A：FY2029 起 AI 營益率額外變動", "pt", None, 0.0, None, "Assumed", "方向：A 最有利", "—")
inp("dm_B", "情境 B：FY2029 起 AI 營益率額外變動", "pt", None, -0.01, None, "Assumed", "需求放緩時價格競爭（Supermicro、HPE、ODM 直供）", "—")
inp("dm_C", "情境 C：FY2029 起 AI 營益率額外變動", "pt", None, -0.025, None, "Assumed", "同上，更深", "—")
inp("m_naif", "非 AI ISG 營業利益率（FY2028 起）", "%", 0.135, 0.155, 0.175, "Assumed", "FY2027 傳統伺服器受記憶體漲價帶動（2Q 年增 122%），FY2028 起正常化", "—")
inp("m_csgf", "CSG 營業利益率（FY2028 起）", "%", 0.06, 0.07, 0.08, "Assumed", "FY26 5.6%、1H FY27 7.8%（公司）", "記憶體成本上升為下檔")

sec("E. 非 AI 營收成長（不隨 AI 情境）")
for y, (lo, b, hi) in ((28, (-0.15, -0.05, 0.05)), (29, (-0.05, 0.02, 0.06)), (30, (-0.03, 0.03, 0.06))):
    inp(f"g_tr{y}", f"傳統伺服器與網通年增 FY20{y}", "%", lo, b, hi, "Assumed", "FY27 1H 年增約 115%（記憶體轉嫁與換機潮，公司）" if y == 28 else "同上", "—")
for y, (lo, b, hi) in ((28, (-0.03, 0.04, 0.10)), (29, (-0.02, 0.04, 0.08)), (30, (-0.02, 0.04, 0.08))):
    inp(f"g_st{y}", f"儲存年增 FY20{y}", "%", lo, b, hi, "Assumed", "FY26 +1%、2Q FY27 +26%（公司）" if y == 28 else "同上", "—")
for y, (lo, b, hi) in ((28, (-0.06, 0.03, 0.08)), (29, (-0.04, 0.03, 0.06)), (30, (-0.04, 0.03, 0.06))):
    inp(f"g_csg{y}", f"CSG（PC）年增 FY20{y}", "%", lo, b, hi, "Assumed", "FY26 +5%、2Q FY27 +20%（Windows 10 換機、商用；公司）" if y == 28 else "同上", "—")

sec("F. 淨利係數、股數（FY2028 起）")
inp("nio_f", "non-GAAP 淨利 ÷ 分部營業利益", "%", 0.70, 0.75, 0.78, "Derived", "1H FY27 7.781 ÷ 10.148＝0.767；FY26 7.046 ÷ 9.944＝0.709（公司）", "利息費用、稅率（non-GAAP 固定稅率；數值未找到）")
inp("sh_g", "稀釋股數年變動", "%", -0.05, -0.03, -0.01, "Assumed", "FY26→1H FY27：6.84 → 6.54 億股（−4%，公司）", "回購授權與自由現金流（1H FY27 調整後 113 億美元）")

sec("G. 共識（只作 Recon 對帳，不作校準）")
inp("c_rev27", "共識 FY2027 營收（22 位）", "US$B", None, 193.4, None, "Interested-party", SAF + "（2026-10-21）", "—")
inp("c_eps27", "共識 FY2027 non-GAAP EPS（22 位）", "US$", None, 25.90, None, "Interested-party", SAF + "（2026-10-21；高 28.93、低 25.50）", "—")
inp("c_rev28", "共識 FY2028 營收", "US$B", None, 229.12, None, "Interested-party", SAF + "：Revenue Next Year +18.44%", "—")
inp("c_eps28", "共識 FY2028 non-GAAP EPS", "US$", None, 29.93, None, "Interested-party", SAF + "：EPS Next Year +15.56%", "FY2029 共識 EPS：未找到")

sec("H. 價格與市場倍數")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 456.44, None, "Interested-party", SA + "：10-23 收盤 456.44", MS + "：10-23 收盤 456.44")
inp("pe_ttm", "近 12 月 P/E（GAAP，stockanalysis）", "x", None, 26.63, None, "Interested-party", SA, "—")
inp("fpe", "前瞻 P/E（stockanalysis，10-23）", "x", None, 16.22, None, "Interested-party", SA, "—")
inp("cons_feps", "共識前瞻 12 月 EPS（現價 ÷ 前瞻 P/E）", "US$", None, "=px/fpe", None, "Derived", "stockanalysis 前瞻 P/E 16.22（2026-10-23）", "共識前瞻 EPS 的直接數字：未找到")
inp("cons_tp", "市場共識目標價（賣方平均）", "US$", 480, 593, 735, "Interested-party", SAF + "：29 位平均 593、中位數 600、最高 735、最低 480（2026-10-21）", MS + "：29 位平均 593；" + ALP + "：599.25；" + INV + "：23 位 502.78（日期不明）")
inp("so", "流通股數", "十億股", None, 0.63581, None, "Interested-party", SA + "（2026-10-23）", "—")

sec("I. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%；美、台、日、韓同一值）", "淨負債約 237 億美元、股東權益為負，財務槓桿高於同業；未改折現率，只在方法欄說明")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "DCF 永續成長率；FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均）", "%", None, 0.20, None, "Assumed", "股東權益為負（FY26 年底 −24.70 億、2Q FY27 −14.27 億美元，公司），ROE 無意義：未找到", "依工作單用預設 0.20；stockanalysis ROE 顯示 n/a、ROIC 29.36%")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（Andy 2026-10-25 11:53 定各半）", "—")
inp("nm10_k", "十年平均淨利率替代：最後一年淨利率 × 係數", "x", None, 0.8, None, "Assumed", "non-GAAP 十年淨利率序列：未找到（FY2017–FY2020 含 EMC 併購虧損，GAAP 不可比）", "工作單：找不到十年平均用最後一年 × 0.8")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2031 起由 FY2030 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2030 淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2030（公式於建置時連結）", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（FY2030 與替代十年平均的中點）", "%", None, 0, None, "Derived", "PL 頁情境 B FY2030 × (1＋nm10_k) ÷ 2", "—")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（替代十年平均＝FY2030 × nm10_k）", "%", None, 0, None, "Assumed", "PL 頁情境 C FY2030 × nm10_k", "—")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "nm10_k", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    for u_, nm_ in URLNAME.items():
        s1 = s1.replace(u_, nm_); s2 = s2.replace(u_, nm_)
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for k in TP_KEYS:
    for c in range(1, 10): IN.cell(INROW[k], c).fill = KEY
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2027 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2027 錨點（十億美元；EPS 美元；股數十億股；non-GAAP）。藍字＝輸入；黑字＝公式。實績取自 Dell 8-K 財報新聞稿（SEC）"
Q["A2"] = "FY2027 錨點＝公司全年指引（營收 1,920 億、AI 伺服器 740 億、non-GAAP EPS 25.50）＋1H FY27 實績＋2H 假設（儲存、CSG、ISG 與 CSG 營益率、股數）；傳統伺服器與網通為全年營收殘差。分部營業利益合計約等於 non-GAAP 營業利益（FY26 差 0.05、1H FY27 差 0.02）"
hdr(Q, 3, ["項目", "FY2026A", "1H FY27（實績）", "2H FY27（模型）", "FY2027 錨點", "", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2026：Dell 8-K 4Q FY26（2026-02-26）；1H FY27：Dell 8-K 2Q FY27（2026-09-01）"
S2 = "1Q FY27：Dell 8-K 1Q FY27；Pulse 2.0、StorageNewsletter 轉述 2Q 數字相同"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"E{r}"] = ann or f"=C{r}+D{r}"; Q[f"E{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "AI 伺服器（AI-Optimized Servers）", {"B": 24.683, "C": 32.533, "D": "=g_ai-C4"}, ann="=g_ai", tag="Verified／Interested-party（全年指引）", s2="1Q FY27 161.32 億、2Q 164.01 億")
qrow(5, "傳統伺服器與網通", {"B": 19.512, "C": 19.074, "D": "=E5-C5"}, ann="=g_rev-E4-E6-E8-E9", tag="Verified／Derived（全年殘差）", s2="1Q 85.43 億、2Q 105.31 億")
qrow(6, "儲存", {"B": 16.631, "C": 9.184, "D": "=C6*st_h2"}, tag="Verified／Assumed（2H）", s2="1Q 43.34 億、2Q 48.50 億")
qrow(7, "ISG 營收", {"B": "=B4+B5+B6", "C": "=C4+C5+C6", "D": "=D4+D5+D6"}, tag="Derived", s2="FY26 608.26、1H FY27 607.91（公司合計）")
qrow(8, "CSG 營收（商用＋消費）", {"B": 50.984, "C": 29.643, "D": "=C8*(1+csg_h2)"}, tag="Verified／Assumed（2H）", s2="FY26 商用 440.62、消費 69.22；1H FY27 商用 262.12、消費 34.31")
qrow(9, "其他業務（營收 − ISG − CSG）", {"B": "=B10-B7-B8", "C": "=C10-C7-C8", "D": "=C9"}, tag="Derived／Assumed（2H＝1H）", s2="Secureworks 出售後的剩餘業務、VMware 轉售等；營業利益併入淨利係數")
qrow(10, "營收", {"B": 113.538, "C": 90.813, "D": "=E10-C10"}, ann="=g_rev", tag="Verified／Interested-party（全年指引）", s2="FY27 錨點＝全年指引 1,920 億")
qrow(11, "ISG 營業利益", {"B": 7.111, "C": 7.836, "D": "=D7*isg_h2m"}, tag="Verified／Assumed（2H）", s2="2Q FY27 47.81 億（營益率 15.0%）")
qrow(12, "CSG 營業利益", {"B": 2.833, "C": 2.312, "D": "=D8*csg_h2m"}, tag="Verified／Assumed（2H）", s2="2Q FY27 11.42 億（7.6%）")
qrow(13, "分部營業利益合計", {c: f"={c}11+{c}12" for c in "BCD"}, tag="Derived", s2="non-GAAP 營業利益：FY26 99.91、1H FY27 101.64（公司）")
qrow(14, "  non-GAAP 營業利益（公司揭露）", {"B": 9.991, "C": 10.164}, ann="", tag="Verified", s2="—")
qrow(15, "non-GAAP 淨利", {"B": 7.046, "C": 7.781, "D": "=E15-C15"}, ann="=g_eps*E17", tag="Verified／Interested-party（全年＝指引 EPS × 股數）", s2="GAAP 淨利 FY26 59.36、1H FY27 75.71")
qrow(16, "淨利 ÷ 分部營業利益", {c: f"={c}15/{c}13" for c in "BCDE"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2028 起 nio_f 的對照")
qrow(17, "稀釋股數（十億股）", {"B": "=7.046/10.30", "C": 0.654, "D": "=h2_sh"}, ann="=(C17+D17)/2", tag="Verified／Derived／Assumed", s2="FY26＝non-GAAP 淨利 ÷ non-GAAP EPS；1H FY27 公司揭露 6.54 億")
qrow(18, "non-GAAP EPS", {c: f"={c}15/{c}17" for c in "BCD"}, fmt="0.00", ann="=E15/E17", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(19, "  公司揭露 non-GAAP EPS", {"B": 10.30, "C": 11.90}, fmt="0.00", ann="=g_eps", tag="Verified／Interested-party（指引）", s2="GAAP EPS：FY26 8.68、1H FY27 11.58、FY27 指引 24.37")
Q.cell(21, 1, "錨點與拆分的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(22, 1, "FY2027 營收年增"); Q["B22"] = "=E10/B10-1"
Q.cell(23, 1, "2H FY27 傳統伺服器與網通 ÷ 1H − 1"); Q["B23"] = "=D5/C5-1"
Q.cell(24, 1, "FY2026 AI 伺服器營業利益率（隱含；非 AI ISG 取 m_nai）"); Q["B24"] = "=(B11-m_nai*(B5+B6))/B4"
Q.cell(25, 1, "FY2027 AI 伺服器營業利益率（隱含；PL 校準值）"); Q["B25"] = "=(E11-m_nai*(E5+E6))/E4"
Q.cell(26, 1, "FY2027 EPS ÷ 共識（25.90）− 1"); Q["B26"] = "=E18/c_eps27-1"
Q.cell(27, 1, "1H FY27 EPS 自算 − 公司揭露（應約為 0）"); Q["B27"] = "=C18-C19"
Q.cell(28, 1, "2H FY27 non-GAAP EPS（隱含；3Q 指引 6.50）"); Q["B28"] = "=D18"
for rr in range(22, 29): Q[f"B{rr}"].number_format = "0.0%" if rr in (22, 23, 24, 25, 26) else "0.00"
for nm, rr in (("ai", 4), ("tr", 5), ("st", 6), ("csg", 8), ("oth", 9), ("rev", 10), ("isgoi", 11), ("csgoi", 12), ("oi", 13), ("ni", 15), ("sh", 17), ("eps", 18)):
    defname(f"a26_{nm}", f"Q!$B${rr}"); defname(f"a27_{nm}", f"Q!$E${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDE": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 70; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2026, 2027, 2028, 2029, 2030); FC = {2026: "C", 2027: "D", 2028: "E", 2029: "F", 2030: "G", 2031: "H", 2032: "I"}
FADE = (2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
BU = (2027, 2028, 2029, 2030); F3 = (2028, 2029, 2030)
cy = lambda y: (y - 1) % 100          # FY(t) 取 CY(t−1)
DEFP = dict(scn="A", rel="rel_{y}", mf="1", vg="v_gb", vv="v_vr", aidm="ai_dm", mn="m_naif", mc="m_csgf", tr="g_tr{y}", st="g_st{y}", cg="g_csg{y}",
            nio="nio_f", shg="sh_g", mnai="m_nai")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    p = lambda k, y=None: P[k].format(y=(y % 100 if y else ""))
    um = lambda y: {2027: "", 2028: "", 2029: f"*um28_{s}", 2030: f"*um29_{s}"}[y]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "FY2030", "FY2031（收斂）", "FY2032（收斂）", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 10, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. 商用 GPU 系統市場與 Dell 份額（D 欄＝FY2027 校準年≈CY2026；E–G＝FY2028–FY2030≈CY2027–29）")
    mfx = lambda y: "" if y == 2027 else f"*{P['mf']}"
    row("ug", "GB300 級 GPU 顆數 × 情境係數", "百萬顆", {c_(y): f"=ug_{cy(y)}{um(y)}{mfx(y)}" for y in BU}, fmt="0.00")
    row("uv", "VR200 級 GPU 顆數 × 情境係數", "百萬顆", {c_(y): f"=uv_{cy(y)}{um(y)}{mfx(y)}" for y in BU}, fmt="0.00")
    row("mkt", "商用 GPU 系統市場（顆數 × 每顆系統價值）", "US$B", {c_(y): f"={c_(y)}{R['ug']}*{P['vg']}+{c_(y)}{R['uv']}*{P['vv']}" for y in BU}, fmt="#,##0.0",
        note="百萬顆 × 千美元／顆＝十億美元；ASIC（TPU、Trainium 等）不含")
    row("sh0", "Dell 份額（FY2027 校準）＝ AI 伺服器營收指引 ÷ 市場", "%", {"D": f"=a27_ai/D{R['mkt']}"}, fmt="0.00%",
        note="Derived：分子為公司 FY27 AI 伺服器營收指引；分母為 CY2026 全部商用 GPU 系統價值")
    row("rel", "份額相對值（÷ FY2027；全部情境）", "x", {"D": "=1", **{c_(y): f"={p('rel', y)}" for y in F3}}, fmt="0.00")
    row("srel", f"情境 {s} 份額乘數（FY2029 起）", "x", {"D": "=1", "E": "=1", "F": f"=srel_{s}", "G": f"=srel_{s}"}, fmt="0.00")
    row("sh", "Dell 有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}*{c_(y)}{R['srel']}" for y in BU}, fmt="0.00%")
    row("ai", "AI 伺服器營收", "US$B", {"C": "=a26_ai", **{c_(y): f"={c_(y)}{R['mkt']}*{c_(y)}{R['sh']}" for y in BU}}, fmt="#,##0.0")
    sec_("B. 非 AI 營收（傳統伺服器、儲存、PC；不隨 AI 情境）")
    row("tr", "傳統伺服器與網通", "US$B", {"C": "=a26_tr", "D": "=a27_tr", **{c_(y): f"={c_(y - 1)}{r}*(1+{p('tr', y)})" for y in F3}}, fmt="#,##0.0")
    row("st", "儲存", "US$B", {"C": "=a26_st", "D": "=a27_st", **{c_(y): f"={c_(y - 1)}{r}*(1+{p('st', y)})" for y in F3}}, fmt="#,##0.0")
    row("isg", "ISG 營收", "US$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['tr']}+{c_(y)}{R['st']}" for y in FYS}, fmt="#,##0.0")
    row("csg", "CSG 營收（PC）", "US$B", {"C": "=a26_csg", "D": "=a27_csg", **{c_(y): f"={c_(y - 1)}{r}*(1+{p('cg', y)})" for y in F3}}, fmt="#,##0.0")
    row("oth", "其他業務", "US$B", {"C": "=a26_oth", "D": "=a27_oth", **{c_(y): f"={c_(y - 1)}{r}" for y in F3}}, fmt="#,##0.0", note="FY2028 起持平（Assumed）")
    row("rev", "營收", "US$B", {c_(y): f"={c_(y)}{R['isg']}+{c_(y)}{R['csg']}+{c_(y)}{R['oth']}" for y in FYS}, fmt="#,##0.0")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI 伺服器 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    sec_("C. 損益（non-GAAP）")
    row("aim", "AI 伺服器營業利益率", "%", {"C": f"=(a26_isgoi-{P['mnai']}*(a26_tr+a26_st))/a26_ai", "D": f"=(a27_isgoi-{P['mnai']}*(a27_tr+a27_st))/a27_ai",
                                             "E": f"=$D{r}+{P['aidm']}", "F": f"=$D{r}+{P['aidm']}+dm_{s}", "G": f"=$D{r}+{P['aidm']}+dm_{s}"}, fmt="0.00%",
        note="FY2026、FY2027 為殘差（ISG 營業利益 − 非 AI 營益率 m_nai × 非 AI 營收）；FY2028 起＝FY2027＋ai_dm（FY2029 起再加情境變動）")
    row("naim", "非 AI ISG 營業利益率", "%", {"C": f"={P['mnai']}", "D": f"={P['mnai']}", **{c_(y): f"={P['mn']}" for y in F3}}, fmt="0.00%")
    row("isgoi", "ISG 營業利益", "US$B", {"C": "=a26_isgoi", "D": "=a27_isgoi", **{c_(y): f"={c_(y)}{R['ai']}*{c_(y)}{R['aim']}+({c_(y)}{R['tr']}+{c_(y)}{R['st']})*{c_(y)}{R['naim']}" for y in F3}}, fmt="#,##0.00")
    row("csgm", "CSG 營業利益率", "%", {"C": f"=a26_csgoi/C{R['csg']}", "D": f"=a27_csgoi/D{R['csg']}", **{c_(y): f"={P['mc']}" for y in F3}}, fmt="0.00%")
    row("csgoi", "CSG 營業利益", "US$B", {c_(y): f"={c_(y)}{R['csg']}*{c_(y)}{R['csgm']}" for y in FYS}, fmt="#,##0.00")
    row("oi", "分部營業利益（≈ non-GAAP 營業利益）", "US$B", {c_(y): f"={c_(y)}{R['isgoi']}+{c_(y)}{R['csgoi']}" for y in FYS}, fmt="#,##0.00")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("nio", "淨利 ÷ 營業利益", "%", {"C": "=a26_ni/a26_oi", "D": "=a27_ni/a27_oi", **{c_(y): f"={P['nio']}" for y in F3}}, fmt="0.0%")
    row("ni", "non-GAAP 淨利", "US$B", {"C": "=a26_ni", "D": "=a27_ni", **{c_(y): f"={c_(y)}{R['oi']}*{c_(y)}{R['nio']}" for y in F3}}, fmt="#,##0.00")
    row("sh_", "稀釋股數", "十億股", {"C": "=a26_sh", "D": "=a27_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['shg']})" for y in F3}}, fmt="0.0000")
    row("eps", "non-GAAP EPS", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("nim", "淨利率（non-GAAP 淨利 ÷ 營收）", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("pgpu", "Dell AI 營收 ÷ 商用 GPU 顆數（每顆）", "US$K/顆", {c_(y): f"={c_(y)}{R['ai']}/({c_(y)}{R['ug']}+{c_(y)}{R['uv']})" for y in BU}, fmt="0.000")
    # FY2031–32：收斂法（上游只到 CY2029＝FY2030；非驅動因子推導）
    for y in FADE:
        c, pc, k = FC[y], FC[y - 1], y - 2030
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+(fade_g{s}-$G{R['g']})*{k}/2"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim']}+(fade_m{s}-$G{R['nim']})*{k}/2"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['sh_']}"] = f"={pc}{R['sh_']}*(1+{P['shg']})"
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh_']}"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("g", "0.0%"), ("rev", "#,##0.0"), ("nim", "0.00%"), ("ni", "#,##0.00"), ("sh_", "0.0000"), ("eps", "0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fm; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 10, "FY2031 起為收斂法：營收年增由 FY2030 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Dell 損益（十億美元；EPS 美元；non-GAAP）。FY2026A 實績、FY2027 錨點（公司指引＋1H 實績＋2H 假設；三情境相同），FY2028 起為模型。FY(t)≈CY(t−1)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2027–FY2028 三情境相同（顆數係數從 CY2028＝FY2029 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數係數 1.0", "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85；份額 ×0.95、AI 營益率 −1pt", "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6；份額 ×0.85、AI 營益率 −2.5pt"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "FY2030", "FY2031（收斂）", "FY2032（收斂）", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "US$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 伺服器營收", "US$B", "#,##0.0"),
        ("tr", "  傳統伺服器與網通", "US$B", "#,##0.0"), ("st", "  儲存", "US$B", "#,##0.0"), ("csg", "  CSG（PC）", "US$B", "#,##0.0"),
        ("aim", "AI 伺服器營業利益率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"), ("ni", "non-GAAP 淨利", "US$B", "#,##0.00"), ("nim", "  淨利率", "%", "0.00%"),
        ("eps", "non-GAAP EPS", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg") and y == 2026: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{SUMR[('A', 'nim')]}")
IN.cell(INROW["fade_mB"], 5, f"=PL!G{SUMR[('B', 'nim')]}*(1+nm10_k)/2")
IN.cell(INROW["fade_mC"], 5, f"=PL!G{SUMR[('C', 'nim')]}*nm10_k")
for k in ("fade_mA", "fade_mB", "fade_mC"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2027 模型 − 錨點").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（US$B）", "+".join(f"ABS(D{PB[s]['rev']}-a27_rev)" for s in SC)),
               ("  三情境：EPS 差（US$）", "+".join(f"ABS(D{PB[s]['eps']}-a27_eps)" for s in SC)),
               ("  情境 A：ISG 營業利益差（以 FY2028 公式套用 FY2027）", f"D{PB['A']['ai']}*D{PB['A']['aim']}+(D{PB['A']['tr']}+D{PB['A']['st']})*D{PB['A']['naim']}-a27_isgoi")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 60, sr
PL.cell(2, 1, "灰底＝FY2031–FY2032 收斂法（上游顆數只到 CY2029＝FY2030）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 I 節），非驅動因子推導")
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["J"].width = 70
for c in "CDEFGHI": PL.column_dimensions[c].width = 13

# ---------------- Units ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 商用 GPU 系統市場（tsmc 顆數 × 橋接表每顆系統價值）、Dell 份額與組裝廠份額對照（情境 A）。FY(t)≈CY(t−1)"
A_ = PB["A"]
ucol = lambda y: chr(ord("C") + y - 2027)
hdr(UN, 3, ["項目", "單位", "CY2026（FY2027）", "CY2027（FY2028）", "CY2028（FY2029）", "CY2029（FY2030）", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("ug", "GB300 級 GPU 顆數（Blackwell＋Hopper）", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['ug']}", "0.00", "tsmc 成品"),
        ("uv", "VR200 級 GPU 顆數（Rubin＋Rubin Ultra；2027 起＋AMD）", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['uv']}", "0.00", "tsmc 成品"),
        ("vg", "每顆系統價值 GB300 級", "US$K/顆", lambda y: "=v_gb", "0.0", "橋接表機架價 ÷ 72"),
        ("vv", "每顆系統價值 VR200 級", "US$K/顆", lambda y: "=v_vr", "0.0", "同上"),
        ("mkt", "商用 GPU 系統市場（整櫃售價）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mkt']}", "#,##0.0", "OEM／整櫃售價口徑（含 GPU）"),
        ("sh", "Dell 有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.00%", "FY2027 以 AI 營收指引 740 億校準"),
        ("ai", "Dell AI 伺服器營收", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}", "#,##0.0", ""),
        ("racks", "隱含 Dell 出貨櫃數（NVL72 等值，按當年組合）", "千櫃", lambda y: f"=PL!{FC[y]}{A_['ai']}/PL!{FC[y]}{A_['mkt']}*(PL!{FC[y]}{A_['ug']}+PL!{FC[y]}{A_['uv']})/72*1000", "#,##0.0", "營收 ÷ 平均每櫃售價"),
        ("qmk", "對照：AI 伺服器組裝市場（ODM 價，廣達模型；含 ASIC）", "US$B", None, "#,##0.0", QMK[1]),
        ("dq", "Dell AI 營收 ÷ 組裝市場（ODM 口徑）", "%", lambda y: f"={ucol(y)}{UR['ai']}/{ucol(y)}{UR['qmk']}", "0.00%", "Dell 為 OEM 售價、含加價，同口徑時份額更低"),
        ("wsum", "對照：鴻海＋廣達＋緯創（含緯穎）合計有效份額（緯創模型）", "%", None, "0.00%", WSUM[1]),
        ("tot", "合計：ODM 三家＋Dell（未扣重複）", "%", lambda y: f"={ucol(y)}{UR['wsum']}+{ucol(y)}{UR['dq']}", "0.00%", "Dell 的機架多由鴻海、緯創代工，兩者重複計算；未扣重複仍低於 100%")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit); UR[key] = ur
    for j, y in enumerate(BU):
        if f:
            v = f(y)
        else:
            v = (QMK if key == "qmk" else WSUM)[0][j]
        UN[f"{ucol(y)}{ur}"] = v; UN[f"{ucol(y)}{ur}"].number_format = fmt
        if not f: UN[f"{ucol(y)}{ur}"].font = BLUE
    UN.cell(ur, 7, note); ur += 1
ur += 1
UN.cell(ur, 1, "讀法：Dell 份額的分母是 NVIDIA 與 AMD 商用 GPU 的整櫃售價（不含 Google TPU、Trainium 等 ASIC，Dell 不是其主要供應商）。FY2027 份額約一成，換算到廣達模型的 ODM 組裝市場（含 ASIC）約 12%；鴻海、廣達、緯創（含緯穎）合計約 44%，加上 Dell 仍約 55%，其餘為 Supermicro、HPE、Lenovo、英業達與雲端自建。Dell 的機架多由 ODM 代工，兩者營收有重複，所以實際份額合計更低，不會超過 100%。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 62; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 16
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元；non-GAAP EPS）。目標價＝DCF 與 Comparable 加權（第 0 節）；結論另列「現價 ÷ FY2027、FY2028 EPS（模型）」；十年 P/E 規則值（第 3 節，年底落後 P/E）只作參考"
V["A2"] = "循環位置：Dell 不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照廣達、緯穎）。營收到 FY2030 仍在增加時，見頂年未知，判為中段"
O = 53
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-23 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "#,##0.00"
for k, y, lab in ((3, 2027, "現價 ÷ FY2026 EPS（實績）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    yy = 2026 if k == 3 else y
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[yy]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 B）≈ 3/12 × FY2027＋9/12 × FY2028（2026-11～2027-10）"); V[f"B{O + 7}"] = f"=3/12*PL!D{SUMR[('B', 'eps')]}+9/12*PL!E{SUMR[('B', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 B）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（GAAP，stockanalysis）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis，10-23）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 4}"] = "FY2027 三情境相同（錨點）"; V[f"E{O + 5}"] = "FY2028 三情境相同（顆數係數從 CY2028＝FY2029 開始）"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型營收最高年）").font = BOLD
hdr(V, CY0 + 1, ["情境", "營收最高年（FY2027–FY2030）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2027,2028,2029,2030}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2030,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2030,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（美元；約 FY2027 年底＝2027-01，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!D{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2028 年底（2028-01）價 ＝ 現價 ×（1＋報酬）所需的倍數（÷ FY2029 EPS；約一年後的價格看下一年 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2028 年底價 ÷ 現價", "所需倍數（÷ FY2029 EPS，A）", "所需倍數（÷ FY2029 EPS，B）", "所需倍數（÷ FY2029 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝約 21%。所需倍數低於歷史中位數時，現價相對便宜。歷史倍數為 GAAP 落後 P/E（含一次性稅務與 VMware 利益），與本模型 non-GAAP EPS 口徑不同。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
# 十年 P/E：companiesmarketcap 年底 P/E（GAAP）；Macrotrends 會計年度年底（次年 1 月底）股價 ÷ 近四季 GAAP EPS。2018-12 重新上市（DELL），2016–2018 無可比市價
PE_CMC = [None, None, None, 4.12, 10.5, 5.79, 15.5, 20.1, 19.6, 16.5]
PE_MT = [None, None, None, 3.51, 8.01, 7.35, 11.89, 18.34, 16.35, None]
T10 = vr + 2
V.cell(T10, 1, "5. 十年 P/E（2016–2025 年底；兩來源平均，分位排名用；2018-12 重新上市，2016–2018 無可比市價）").font = BOLD
hdr(V, T10 + 1, ["年底", "兩來源平均 P/E", "companiesmarketcap", "Macrotrends"])
RB0 = T10 + 16
for i, y in enumerate(range(2016, 2026)):
    rr = T10 + 2 + i
    V.cell(rr, 1, str(y))
    if y >= 2019:
        V.cell(rr, 2, f"=AVERAGE(C{rr}:D{rr})").number_format = "0.00"
        V.cell(rr, 3, f"=B{RB0 + 5 + i}").number_format = "0.00"
        V.cell(rr, 4, f'=IF(C{RB0 + 5 + i}="","",C{RB0 + 5 + i})').number_format = "0.00"
PE10 = f"$B${T10 + 2}:$B${T10 + 11}"
V.cell(T10 + 12, 1, "年數"); V[f"B{T10 + 12}"] = f"=COUNT({PE10})"
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；GAAP；2019–2025 共 7 年）"},
    {"label": "Macrotrends（會計年度年底 P/E）", "vals": PE_MT, "tag": "Interested-party（資料商；次年 1 月底股價 ÷ 近四季 GAAP EPS；2019–2024 共 6 年，2025 年（2026-01）未列）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。Dell 2018-12 重新上市，只有 2019–2025。兩來源都是 GAAP P/E：2019–2021 含稅務利益與 VMware 持股利益，倍數偏低（4–8x），規則值偏低。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025", (V.cell(RB0 + 5, 1).value, V.cell(RB0 + 14, 1).value)
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# 0. 目標價（Andy 2026-10-25 11:28／11:53：DCF 與 Comparable 各半；列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期）
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
TT = [1 / 12 + k for k in range(6)]; FR = [1 / 12] + [1] * 5      # FY2027 結束於 2027-01：基準日 2026 年底後只剩一個月
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明", "折現年數 t", "期內占比（基準日後）"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境（non-GAAP）；FY2031 起為收斂法；F 欄＝距 2026 年底的年數（會計年度 1 月底結束）"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "0.00"); EPSR[y] = 8 + i
    V.cell(8 + i, 6, round(TT[i], 6)).number_format = "0.000"; V.cell(8 + i, 7, round(FR[i], 6)).number_format = "0.000"
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
V["E14"] = "下列六列＝EPS × (1 − 再投資率) × 期內占比（FY2027 只計 2027-01 一個月）"
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "tp_roe 用預設 0.20（股東權益為負，ROE 無意義）")
CFR = {}
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)*$G${EPSR[y]}" for c in "BCD"], "0.00"); CFR[y] = 16 + i
def dcf_formula(c, r="tp_r", g="tp_g", reinv="$B$15"):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底；FY2027 只計一個月。"""
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})*Val!$G${EPSR[y]}"
    pv = "+".join(f"{cf(y)}/(1+{r})^Val!$F${EPSR[y]}" for y in YRS6)
    return f"MAX(0,{pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^Val!$F${EPSR[2032]})"
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）",
     [f"=MAX(0,SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^$F$8:$F$13)+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^$F$13)" for c in "BCD"], "0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^t；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底；小於 0 時取 0")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("Hewlett Packard Enterprise", "HPE", 12.19, "2026-10-23"), ("Super Micro Computer（美超微）", "SMCI", 7.06, "2026-10-23"),
         ("Lenovo（聯想）", "0992.HK", 9.60, "2026-10-20"), ("HP Inc.", "HPQ", 11.26, "2026-10-23"), ("NetApp", "NTAP", 21.69, "2026-10-23")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "伺服器、儲存與 PC 品牌同業：HPE、Supermicro、Lenovo、HP Inc.、NetApp（不含 A 股）")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(3/12*B{EPSR[2028]}+9/12*B{EPSR[2029]})+tp_wB*(3/12*C{EPSR[2028]}+9/12*C{EPSR[2029]})+tp_wC*(3/12*D{EPSR[2028]}+9/12*D{EPSR[2029]})", None, None], "0.00",
     "約 2027-10 時的未來 12 月（2027-11～2028-10）EPS ≈ 3／12 × FY2028＋9／12 × FY2029（會計年度 1 月底結束，依月份調整）")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "對照", "對照", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁（0.5）")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "0.00", "stockanalysis：29 位分析師平均 593（中位數 600、最高 735、最低 480），2026-10-21；MarketScreener 29 位 593")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-23 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/C44-1", "=cons_feps", None], "0.0%", "B＝本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；C＝共識前瞻 EPS（現價 ÷ 前瞻 P/E 16.22）；兩者期間差約一年")
V["C44"].number_format = "0.00"
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=C45/D45-1", "=cons_tp/cons_feps", "=B33"], "0.0%")
V["C45"].number_format = "0.0"; V["D45"].number_format = "0.0"
V["E45"] = "C＝共識目標價 ÷ 共識前瞻 EPS 隱含的倍數；D＝同業中位數；B＝C ÷ D − 1"
tset(46, "band", "價格帶（DCF 不如預期～順利）", ["=D23", "=B23", None], "0.00", "B＝C 情境 DCF 價值、C＝A 情境 DCF 價值（皆推到 12 個月後）")
V["C46"].number_format = "0.00"
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2019–2025 年底 P/E（第 5 節兩來源平均，GAAP）的分位；低於最小記 0%、高於最大記 100%；只作檢查，不進目標價")
tset(48, "grade", "上檔空間級距", ['=IF(B42>0.5,"大於 50%",IF(B42>0.25,"25%–50%",IF(B42>=0,"0%–25%",IF(B42>=-0.25,"0%～−25%",IF(B42>=-0.5,"−25%～−50%","小於 −50%")))))', None, None], "@",
     "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
tset(49, "rec", "投資建議", ['=IF(B42>0.5,"強力買進",IF(B42>0.25,"買進",IF(B42>=-0.25,"中立",IF(B42>=-0.5,"放空","強力放空"))))', None, None], "@",
     "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
V.cell(50, 1, "目標價方法")
V["B50"] = ("DCF：股東現金流＝non-GAAP EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 用預設 20%（股東權益為負，近三年 ROE 無意義；再投資率 20%），FY2027–32 明確期＋永續；"
            "基準日 2026 年底，FY2027 只計 2027-01 一個月，之後逐年折現。組裝型 OEM 資本支出低，未改用淨利＋折舊－資本支出。"
            "Comparable：HPE、Supermicro、Lenovo、HP Inc.、NetApp 前瞻 P/E 中位數（不含 A 股）× 2027-11～2028-10 的 EPS（3/12 × FY2028＋9/12 × FY2029）。"
            "偏誤：non-GAAP EPS 排除攤銷與股票薪酬（偏高）；同業多為成長較低的 PC／伺服器公司，未給 Dell 的 AI 份額溢價（Comparable 偏低）；淨負債高、股東權益為負未反映在折現率。"
            "情境對應：A＝順利（顆數基準）、B＝平平（顆數 ×0.8／0.85）、C＝不如預期（顆數 ×0.6）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；Dell 份額、AI 營益率在每個案例內依 FY2027 錨點重算，FY2027 恆等於錨點）：對 FY2028、FY2029 營收、營業利益率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2028 營收", "FY2028 營益率", "FY2028 EPS", "FY2028 EPS 相對基準", "FY2029 營收", "FY2029 營益率", "FY2029 EPS", "FY2029 EPS 相對基準"])
cases = [("AI 部署 ×0.8", "FY2028 起商用 GPU 顆數 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("份額流失（低）", "份額相對值 FY2028 0.75、FY2029 0.65、FY2030 0.55", dict(rel="rel_{y}_lo")),
         ("份額維持（高）", "份額相對值 1.05／1.0／1.0", dict(rel="rel_{y}_hi")),
         ("每顆系統價值 低", "GB300／VR200 機架價取橋接表低值", dict(vg="v_gb_lo", vv="v_vr_lo")), ("每顆系統價值 高", "取高值", dict(vg="v_gb_hi", vv="v_vr_hi")),
         ("AI 營益率 −3pt", "FY2028 起 AI 營益率比 FY2027 低 3 個百分點", dict(aidm="ai_dm_lo")), ("AI 營益率 +1pt", "高 1 個百分點", dict(aidm="ai_dm_hi")),
         ("非 AI 營益率拆分 19%", "FY2027 拆分用 m_nai 19%（AI 營益率校準值較低）", dict(mnai="m_nai_hi")),
         ("非 AI ISG 營益率 低", "FY2028 起 13.5%", dict(mn="m_naif_lo")), ("非 AI ISG 營益率 高", "17.5%", dict(mn="m_naif_hi")),
         ("CSG 營益率 低", "6%", dict(mc="m_csgf_lo")), ("CSG 營益率 高", "8%", dict(mc="m_csgf_hi")),
         ("傳統伺服器 低", "FY2028 −15%、FY2029 −5%、FY2030 −3%", dict(tr="g_tr{y}_lo")), ("傳統伺服器 高", "+5%／+6%／+6%", dict(tr="g_tr{y}_hi")),
         ("PC 成長 低", "FY2028 −6%、之後 −4%", dict(cg="g_csg{y}_lo")), ("PC 成長 高", "+8%／+6%／+6%", dict(cg="g_csg{y}_hi")),
         ("淨利 ÷ 營業利益 0.70", "FY2028 起 0.70", dict(nio="nio_f_lo")), ("回購 少（−1%）", "股數年變動 −1%", dict(shg="sh_g_hi"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2028, "#,##0.0"), ("D", "oim", 2028, "0.00%"), ("E", "eps", 2028, "0.00"),
                           ("G", "rev", 2029, "#,##0.0"), ("H", "oim", 2029, "0.00%"), ("I", "eps", 2029, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"F{i}"] = f"=E{i}/E$4-1"; SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c in "FJ": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 42; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 B", "CY2028 顆數 ×0.8、CY2029 ×0.85；份額 ×0.95、AI 營益率 −1pt", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "CY2028–29 顆數 ×0.6；份額 ×0.85、AI 營益率 −2.5pt", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：情境 B、C 的差異從 FY2029（CY2028）起。FY2028 EPS 主要受 Dell 份額、每顆系統價值與 AI 營益率影響；非 AI（傳統伺服器、儲存、PC）約占營業利益四成，營益率假設影響也大。每顆系統價值的變動大多被 FY2027 份額校準抵銷（只剩世代組合效果）。")
SENS_NOTE = nn + 3
TPS0 = nn + 5
SE.cell(TPS0, 1, "目標價參數（只改 DCF 參數；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "#,##0"
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.3", "", dict(w="0.3")), ("DCF 權重 0.7", "", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
assert TPS_LAST < 42, TPS_LAST
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 48
for c in "CDEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、共識與組裝廠份額。只作事後對照，不為對上數字而調參數"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 共識", "=c_rev27", pl("rev", 2027), "stockanalysis（22 位，10-21）", "錨點＝公司指引 1,920 億", fmt="#,##0.0")
rec("FY2027 EPS vs 共識", "=c_eps27", pl("eps", 2027), "stockanalysis（22 位）", "錨點＝公司指引 25.50")
rec("FY2028 營收 vs 共識", "=c_rev28", pl("rev", 2028), "stockanalysis（Revenue Next Year）", "", fmt="#,##0.0")
rec("FY2028 EPS vs 共識", "=c_eps28", pl("eps", 2028), "stockanalysis（EPS Next Year）", "")
rec("FY2029 EPS vs 共識", None, pl("eps", 2029), "未找到", "共識 FY2029 EPS：未找到")
RC.cell(rr_, 1, "B. 公司指引與在手訂單").font = BOLD; rr_ += 1
rec("FY2027 AI 伺服器營收 vs 指引 740 億", "=g_ai", pl("ai", 2027), "Dell 8-K 2Q FY27", "校準點（差距應為 0）", fmt="#,##0.0")
rec("FY2028 AI 伺服器營收 vs FY2027 指引", "=g_ai", pl("ai", 2028), "Dell 8-K 2Q FY27", "2Q FY27 在手訂單 950 億美元；近四季訂單 1,317 億美元", fmt="#,##0.0")
RC.cell(rr_, 1, "C. 份額與組裝廠合計（廣達、鴻海、緯創模型；唯讀）").font = BOLD; rr_ += 1
rec("FY2027 Dell AI 營收 ÷ 組裝市場（ODM 口徑）", None, lambda s: f"=Units!C{UR['dq']}", "quanta Units", "Dell 為 OEM 售價，同口徑時更低", fmt="0.00%")
rec("FY2027 ODM 三家＋Dell 合計份額（未扣重複）", None, lambda s: f"=Units!C{UR['tot']}", "wistron Units", "應遠低於 100%；其餘為 Supermicro、HPE、Lenovo、英業達、自建", fmt="0.00%")
rec("FY2028 Dell 有效份額", None, pl("sh", 2028), "—", "份額相對值 0.90（Assumed）", fmt="0.00%")
rec("FY2028 Dell AI 營收 ÷ 商用 GPU 顆數（US$K／顆）", None, pl("pgpu", 2028), "—", "", fmt="0.000")
RC.cell(rr_, 1, "D. 校準").font = BOLD; rr_ += 1
rec("FY2027 AI 伺服器營業利益率（Derived）", None, pl("aim", 2027), "FY2027 錨點", "ISG 營業利益 − 16% × 非 AI 營收", fmt="0.00%")
rec("FY2026 AI 伺服器營業利益率（Derived）", None, pl("aim", 2026), "FY2026 實績", "", fmt="0.00%")
rec("2H FY27 non-GAAP EPS（模型）", None, lambda s: "=Q!D18", "3Q FY27 指引 6.50", "2H＝全年指引 − 1H 實績", fmt="0.00")
RC.cell(rr_, 1, "E. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代 tsmc 顆數 × 每顆系統價值"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_LAST = rr_ - 1
RC.column_dimensions["A"].width = 60
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文為 SEC 8-K 附件 99.1；指引為 Interested-party；券商與資料商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：8-K 4Q FY26（FY2026 全年實績、分部、FY27 初始指引；股東權益）", K4),
        ("公司：8-K 1Q FY27（1Q 實績、分部、上修指引；股東權益）", K1),
        ("公司：8-K 2Q FY27（2Q、1H 實績、分部、3Q 與全年指引；股東權益）", K2),
        ("媒體：Pulse 2.0 2026-09-01（AI 訂單、在手訂單、近四季訂單、自由現金流）", PUL),
        ("媒體：StorageNewsletter 2026-09-03（儲存全年約 200 億、3Q 儲存與 ISG 成長）", SNL),
        ("媒體：Zacks 2026-09（non-GAAP 毛利率、指引區間、回購）", ZK),
        ("資料商：stockanalysis 統計（價格、股數、P/E、前瞻 P/E）", SA),
        ("資料商：stockanalysis 預估（共識 EPS、營收、目標價）", SAF),
        ("資料商：MarketScreener（10-23 收盤、目標價 29 位）", MS),
        ("資料商：Alpha Spread（目標價 599.25）", ALP),
        ("資料商：Investing.com（目標價 23 位 502.78）", INV),
        ("資料商：companiesmarketcap 年底 P/E", CMC),
        ("資料商：Macrotrends P/E（會計年度年底）", MT),
        ("資料商：stockanalysis 同業前瞻 P/E（HPE、SMCI、HPQ、NTAP、0992.HK）", "；".join(SAP.format(t) for t in ("hpe", "smci", "hpq", "ntap")) + "；https://stockanalysis.com/quote/hkg/0992/statistics/"),
        ("市場層與組裝廠模型：tsmc、markets/chip_bridge、quanta、wistron 成品（本 repo，唯讀）", f"{rel(TSMC_F)}；{rel(BR_F)}；{rel(QU_F)}；{rel(WI_F)}"),
        ("待補（人工）：3Q FY27 財報（約 2026-11 下旬）；AI 伺服器營益率；non-GAAP 稅率；共識 FY2029 EPS；十年 non-GAAP 淨利率；主要客戶占比（10-K）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
OBS = ["目標價 US$367（DCF 412、Comparable 321 各半），比現價 456.44（2026-10-23）低 19.7%，級距 0%～−25%；Comparable 單獨為 −29.6%。主因是倍數不是 EPS：本模型 12 個月後前瞻 EPS 28.54 只比共識前瞻 EPS 28.14 高 1.4%，但同業（HPE、Supermicro、Lenovo、HP Inc.、NetApp）前瞻 P/E 中位數 11.3 倍，現價隱含 16.2 倍、共識目標價 593 隱含 21.1 倍（比同業高 87%）；市場給 Dell 的 AI 份額溢價，Comparable 未計入。",
       "FY2028 EPS 30.20（三情境相同）比共識 29.93 高 0.9%，前提是 AI 伺服器營收由 FY2027 指引 740 億美元升到 1,221 億美元（+65%；份額相對值 0.90 為 Assumed），2Q FY27 在手訂單 950 億美元可支撐約八成。追蹤 3Q FY27 財報（約 2026-11 下旬）的 AI 訂單、在手訂單與出貨。",
       "AI 伺服器營益率是最大的不確定：公司不揭露，本模型以非 AI ISG 營益率 16% 拆分得 FY2027 約 11.1%（FY2026 約 5.4%）；拆分改用 19% 時 FY2028 EPS −11%，AI 營益率 −3pt 時 −10%。3Q FY27 的 ISG 營益率（2Q 為 15.0%）是第一個檢驗點。"]   # 重點觀察（與 README.md 相同文字）
lines = [("Dell Technologies（DELL US）收支與估值模型 v0.1（AI 半導體 Project：AI 伺服器品牌，樣板＝wiwynn/）　2026-10-25", True),
 ("版本：v0.1（初版）。Dell 是企業與 neocloud 的 AI 伺服器品牌（ISG：AI 伺服器、傳統伺服器與網通、儲存）兼 PC 大廠（CSG）。AI 伺服器以加速器顆數 × 每顆系統價值 × Dell 份額推；傳統伺服器、儲存、PC 另列，不隨 AI 情境。", False),
 ("模型命題：Dell 的 AI 伺服器營收＝商用 GPU 顆數（NVIDIA、AMD；tsmc 成品，同字母情境係數）× 每顆系統價值（橋接表機架價 ÷ 72）× Dell 有效份額（FY2027 以公司 AI 營收指引 740 億美元校準）× 份額相對值；營業利益＝AI 營益率（FY2027 殘差校準）× AI 營收＋非 AI ISG 與 CSG 營益率 × 營收。", False),
 ("核心問題：FY2027 指引 EPS 25.50 美元已比年初指引高一倍；在手訂單 950 億美元能否在 FY2028 轉成 AI 營收再成長、同時守住約 10% 的 AI 營益率？現價約 16 倍前瞻 P/E，高於 OEM 同業約 11 倍，隱含什麼路徑？", False),
 ("重點觀察 1：" + OBS[0], False), ("重點觀察 2：" + OBS[1], False), ("重點觀察 3：" + OBS[2], False),
 ("會計年度：FY(t)＝t−1 年 2 月～t 年 1 月底，FY(t) 直接取 CY(t−1)（FY2027≈CY2026、FY2030≈CY2029）。金額：十億美元。口徑：non-GAAP（公司指引、共識、同業倍數一致）；GAAP EPS 只在 Q 頁對照。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 伺服器營收：tsmc 成品的 Blackwell＋Hopper（GB300 級；2026 另含 AMD MI355）與 Rubin＋Rubin Ultra（VR200 級；2027 起另含 AMD MI450）顆數（世代分組與 Siliconomics Rack 頁、美超微模型相同，CF10） × 同字母情境係數 × 每顆系統價值（橋接表 Tokenomics v20.1 機架價 ÷ 72：GB300 約 59.7、VR200 約 116.7 千美元）＝商用 GPU 系統市場；× Dell 份額（FY2027 校準）× 份額相對值（Assumed）× 情境份額乘數（FY2029 起）。", False),
 ("非 AI：傳統伺服器與網通、儲存、CSG 依成長率推（Assumed，不隨 AI 情境）。營業利益：AI 營益率＝FY2027 殘差（ISG 營業利益 − 非 AI 營益率 16% × 非 AI 營收）＋ai_dm（FY2029 起加情境變動）；非 AI ISG、CSG 營益率為 Assumed。淨利＝分部營業利益 × 淨利係數；股數每年 −3%（回購）。", False),
 ("校準：FY2027 ＝ 公司全年指引（營收、AI 營收、non-GAAP EPS）＋1H FY27 實績＋2H 假設。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），ROE 用預設 20%，FY2027 只計一個月，r 10%、g 4%，三情境權重 0.3／0.5／0.2）＋0.5 × Comparable 目標價（HPE、Supermicro、Lenovo、HP Inc.、NetApp 前瞻 P/E 中位數 × 12 個月後前瞻 EPS；不含 A 股）；另列 Comparable 單獨目標價、市場共識目標價、差距拆解與上檔空間級距。十年 P/E 規則值 × FY2027 EPS 保留為參考（不是目標價）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2028–FY2030 由驅動因子推導；FY2031 起為收斂法，非驅動因子推導（上游顆數到 CY2029＝FY2030；營收年增由 FY2030 線性收斂到 fade_g、淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。", False),
 ("", False),
 ("已知限制：(1) AI 營益率公司不揭露，以非 AI 營益率 16% 拆分（Sens 有 19% 案例）；(2) 份額相對值、情境份額乘數與營益率變動是 Assumed；(3) AMD Helios、Rubin Ultra 每顆系統價值假設同 VR200；(4) 股東權益為負，ROE 用預設 20%；(5) 十年 P/E 只有 2019–2025，且為 GAAP；(6) 共識 FY2029 EPS 未找到；(7) 3Q FY27 財報（約 2026-11 下旬）後更新錨點；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("DELL_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE, "UN_NOTE": UN_NOTE,
               "T": T, "T10": T10, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "PEER_LAST": PEER_LAST, "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0, "NCASES": len(cases)},
              open(os.environ["DELL_ROWMAP"], "w"), ensure_ascii=False, indent=1)
