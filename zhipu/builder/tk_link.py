"""Tokenomics 快照讀取（沿用 OpenAI v0.6 的 63 名；智譜 Z2 另以列標籤讀 Cap_In／Price_Frontier 的 GLM 列與中國合格前緣）。

- 值取自 Tokenomics master 的 model/CURRENT，逐名讀出寫入 TK_Link；不用 Excel 外部連結。
- 列序：P1 原 37 名（列位不變）→ SRC_DEM_010–013（＋013 低／高）→ Block 6（IF_Alloc* 7 名）→ S1 新增（r6 下游需要）→ S4 新增（IF_CapexTotal）→ 智譜 Z3b 新增（IF_CostPre／Cache／Dec_*，9 名）→ 智譜 Z5b 新增（IF_OpexGW、IF_DeprLifeIT）。
- 工作單要求、但 Tokenomics 現行版沒有的名稱：逐名列入、值留空、狀態 PENDING_NOTE（E1；check_tk_snapshot 報 MISSING）。
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

import openpyxl

# 工作單 P1-3 的名稱樣式（依序）
PATTERNS = [
    r"IF_TokGW_(Luna|Sol|Astra)", r"IF_Util", r"IF_HoldEcon", r"IF_FullCost_(Luna|Sol|Astra)",
    r"IF_Price(Fresh|Cached|Think|Out|Ref)_(Luna|Sol|Astra)", r"IF_FrontRef_(Luna|Sol|Astra)",
    r"IF_ProgGWyr_(Luna|Sol|Astra)", r"IF_RDMult", r"IF_TrainGenDefault",
]
SRC_DEM = [f"SRC_DEM_{i:03d}" for i in range(4, 14)] + ["SRC_DEM_013_Lo", "SRC_DEM_013_Hi"]   # 013 只有低／高（無點值）
# Block 6（v5.15 起）：工作單暫定名與 Tokenomics 實際名稱相同；IF_AllocQ1_R2 為 r2 補列
BLOCK6 = ["IF_AllocQ1", "IF_AllocQ1_R2", "IF_AllocQ2", "IF_AllocServeGW", "IF_AllocRDGW", "IF_AllocDemand", "IF_AllocImpliedNk"]
# S1 新增（工程類；r6 下游會用到）：訓練成本（P3 研發檢查）、每 GW 理論營收（P4 命題對照）、
# 全成本下游預設（Tokenomics 標為「下游預設」；P4 對照）、L1 問 3 研發算力成本（P4 Q3，r6 P4-2）
EXTRA_PATTERNS = [
    r"IF_TrainCost_(Luna|Sol|Astra)", r"IF_RevGW_(Luna|Sol|Astra)", r"IF_FullCostDefault_(Luna|Sol|Astra)",
]
EXTRA_NAMES = ["L1_Ans3", "L1_Ans3_Lo", "L1_Ans3_Hi"]
# S4 新增（工程類；P4 自建 GW 需要每 GW 資本支出；列於最後，既有列位不變）
S4_NAMES = ["IF_CapexTotal"]
# 智譜 Z3b（F1）新增：逐 token 類型單位成本（經濟口徑、100% 利用率；新鮮 prefill／快取命中 prefill／decode），列於最後
Z3B_PATTERNS = [r"IF_Cost(Pre|Cache|Dec)_(Luna|Sol|Astra)"]
# 智譜 Z5b（V2）新增：自有算力每 GW 營運費用（不含折舊）與 IT 折舊年限，列於最後
Z5B_NAMES = ["IF_OpexGW", "IF_DeprLifeIT"]
PENDING_NOTE = "待 Tokenomics 提供"
TABLE_NOTE = "讀表（非具名）"

# 智譜 Z2 步驟 2：以列標籤讀表（Tokenomics 無具名範圍者）；只供 Revenue 對照列與 Checks，不作驅動。
# (本模型名稱, 頁, 表頭判定(欄A值, 欄B值或 None), 列標籤（欄 A 完全相符）, 起始欄, 欄數, 說明, 單位)
TABLE_ROWS = [
    ("RD_CapIn_GLM53", "Cap_In", ("模型", "廠商"), "GLM-5.3", "E", 5, "Cap_In F 表 GLM-5.3：新鮮輸入、快取輸入、輸出（國際站 $/M）、尖峰離峰、能力指數", "$/M；指數"),
    ("RD_CapIn_GLM53Flash", "Cap_In", ("模型", "廠商"), "GLM-5.3-Flash", "E", 5, "Cap_In F 表 GLM-5.3-Flash：同上", "$/M；指數"),
    ("RD_PF_GLM53_Mix", "Price_Frontier", ("模型", "國別"), "GLM-5.3", "H", 3, "Price_Frontier A 表 GLM-5.3：參考請求混合單價（Luna、Sol、Astra 參考請求）", "$/M"),
    ("RD_PF_GLM53Flash_Mix", "Price_Frontier", ("模型", "國別"), "GLM-5.3-Flash", "H", 3, "Price_Frontier A 表 GLM-5.3-Flash：同上", "$/M"),
    ("RD_PF_ChinaFront", "Price_Frontier", ("項目", "單位"), "中國廠商 合格者最低 混合單價", "C", 3, "Price_Frontier B 表：中國合格前緣混合單價（Luna、Sol、Astra；無合格者為文字）", "$/M"),
    ("RD_PF_ChinaFrontModel", "Price_Frontier", ("項目", "單位"), "中國廠商 合格者最低 模型", "C", 3, "Price_Frontier B 表：中國合格前緣模型名（Luna、Sol、Astra）", "文字"),
    ("RD_PF_ChinaVsOAI", "Price_Frontier", ("項目", "單位"), "中國合格者最低 ÷ OpenAI 層級模型", "C", 3, "Price_Frontier B 表：中國合格前緣 ÷ OpenAI 層級模型（Luna、Sol、Astra）", "x"),
]


def read_table_rows(wb):
    """依列標籤讀 TABLE_ROWS；找不到標籤時該列狀態為 MISSING（值留空）。"""
    from openpyxl.utils import column_index_from_string as ci
    out = []
    for name, sheet, hdr, label, c0, n, desc, unit in TABLE_ROWS:
        ws = wb[sheet]
        start = None
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == hdr[0] and (hdr[1] is None or ws.cell(r, 2).value == hdr[1]):
                start = r
                break
        row = None
        if start:
            for r in range(start + 1, ws.max_row + 1):
                if str(ws.cell(r, 1).value).strip() == label:
                    row = r
                    break
        vals = [ws.cell(row, ci(c0) + k).value for k in range(n)] if row else []
        out.append(dict(name=name, kind="RD", label=f"{desc}（{sheet}!{c0}{row}，標籤『{label}』）", unit=unit, values=vals,
                        status=TABLE_NOTE if row else "MISSING", sheet=sheet, row=row))
    return out


def _split(attr: str):
    m = re.match(r"^(?:'([^']+)'|([^!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", attr)
    sheet = m.group(1) or m.group(2)
    return sheet, m.group(3), int(m.group(4)), m.group(5), (int(m.group(6)) if m.group(6) else None)


def read_snapshot(tk_dir: Path):
    tk_dir = Path(tk_dir)
    current = (tk_dir / "model" / "CURRENT").read_text(encoding="utf-8").strip()
    path = tk_dir / "model" / current
    sha = subprocess.run(["git", "-C", str(tk_dir), "rev-parse", "HEAD"], capture_output=True, text=True).stdout.strip()
    m = re.search(r"_v(\d+(?:\.\d+)?)\.xlsx$", current)
    wb = openpyxl.load_workbook(path, data_only=True)
    names = {k: v.attr_text for k, v in wb.defined_names.items()}
    chosen = []
    for pat in PATTERNS:
        chosen += sorted(n for n in names if re.fullmatch(pat, n))
    chosen += [n for n in SRC_DEM if n in names]
    chosen += [n for n in BLOCK6 if n in names]
    for pat in EXTRA_PATTERNS:
        chosen += sorted(n for n in names if re.fullmatch(pat, n))
    chosen += [n for n in EXTRA_NAMES if n in names]
    chosen += [n for n in S4_NAMES if n in names]
    for pat in Z3B_PATTERNS:
        chosen += sorted(n for n in names if re.fullmatch(pat, n))
    chosen += [n for n in Z5B_NAMES if n in names]
    rows, seen = [], set()
    for n in chosen:
        if n in seen:
            continue
        seen.add(n)
        sheet, c1, r1, c2, r2 = _split(names[n])
        ws = wb[sheet]
        if c2:
            from openpyxl.utils import column_index_from_string as ci
            vals = [ws.cell(r1, k).value for k in range(ci(c1), ci(c2) + 1)]
        else:
            vals = [ws[f"{c1}{r1}"].value]
        if sheet == "Interface":
            label = re.sub(r"[\s　]*\[IF_[^\]]*\]\s*$", "", str(ws.cell(r1, 1).value))
            unit = ws.cell(r1, 2).value or ""
        elif sheet == "L1":
            col = {"D": "基準", "E": "低", "F": "高"}.get(c1, c1)
            label = f"{ws.cell(r1, 2).value}（{col}）"
            unit = ws.cell(r1, 7).value or ""
        else:
            label = f"{ws.cell(r1, 2).value}（{ws.cell(r1, 7).value}）"
            unit = ws.cell(r1, 6).value or ""
        rows.append(dict(name=n, kind=n.split("_")[0], label=label, unit=unit, values=vals, status="OK"))
    present = {r["name"] for r in rows}
    # 工作單要求、但 Tokenomics 現行版沒有者：逐名列入、值留空（E1）
    wanted = SRC_DEM + BLOCK6 + EXTRA_NAMES + S4_NAMES + [f"IF_Cost{k}_{t}" for k in ("Pre", "Cache", "Dec") for t in ("Luna", "Sol", "Astra")] + Z5B_NAMES
    pending = [dict(name=n, kind=n.split("_")[0], label="", unit="", values=[], status=PENDING_NOTE) for n in wanted if n not in present]
    hdr_gen = [wb["Interface"].cell(4, k).value for k in range(3, 18)]
    hdr_cost = [wb["Interface"].cell(5, k).value for k in range(3, 18)]
    return dict(file=current, version=f"v{m.group(1)}" if m else "?", sha=sha, rows=rows, pending=pending,
                hdr_gen=hdr_gen, hdr_cost=hdr_cost, table=read_table_rows(wb))
