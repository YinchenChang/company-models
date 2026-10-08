#!/usr/bin/env python3
"""產生智譜收支模型 Excel（v0.1-Z3：README、SRC_ZP、TK_Link、OAI_Link、Inputs、Demand、Revenue、Compute、Cost、Checks）。

用法（在 zhipu/ 內）：
  python3 builder/build.py --tk-dir <Tokenomics checkout> --oai-xlsx <OpenAI v0.6 xlsx> --out model/20261008_Zhipu_v0.1.xlsx \
      [--base <前一版>] [--date 2026-10-08] [--oai-ref claude/openai-s6-release@e91df57]
  --base：前一版 xlsx；依 Excel 優先規則保留 Excel 內已改過的藍字輸入（builder/preserve.py）。
輸出經 LibreOffice 重算後才寫到 --out（內含快取值；engine 仍會強制重算）。紀錄檔寫在 build_out/（restore_log.txt、build_summary.json）。

Python 只產生結構：每格公式文字在此組裝；所有數值來自 data/zhipu_src.yaml（SRC_ZP）、data/zhipu_inputs.yaml（Inputs）、
Tokenomics 快照（TK_Link）與 OpenAI v0.6 快照（OAI_Link）。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import preserve  # noqa: E402
import source_rules  # noqa: E402
from common import F_BOLD, F_CALC, F_IN, F_NOTE, header, lo_recalc, put, title  # noqa: E402
from oai_link import read_oai  # noqa: E402
from tk_link import PENDING_NOTE, TABLE_NOTE, read_snapshot  # noqa: E402
import z2  # noqa: E402
import z3  # noqa: E402

VERSION = "v0.1-Z3"
SRC_YAML = ROOT / "data" / "zhipu_src.yaml"
INP_YAML = ROOT / "data" / "zhipu_inputs.yaml"
REGISTRY_PATH = HERE / "id_registry.json"
SRC = preserve.SRC_SHEET
SRC_COLS = ["SRC_ID", "鍵（key）", "項目", "值", "低", "高", "單位", "期間", "來源", "網址", "文件日期", "擷取日", "標記", "利害方與誘因", "口徑（basis）",
            "來源等級", "立場", "立場說明", "一手／二手", "備註", "無數值（公式）", "缺出處（公式）", "區間順序異常（公式）", "四欄標註"]
INP_COLS = ["INP_ID", "鍵（key）", "參數", "索引", "單位", "值", "低", "高", "標記", "依據", "區間理由", "區間順序異常（公式）", "Analogy／Assumed 缺區間（公式）"]
TAGS = ("Verified", "Interested-party", "Analogy", "Assumed", "Derived", "Decision")
TK_FIRST, TK_VCOL0 = 10, 10
OAI_FIRST = 10


def nm(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


# ───────────────────────── 資料與穩定 ID ─────────────────────────
class Data:
    """SRC_ZP 與 Inputs 的列、穩定 ID 與查找函式（S：SRC 鍵→名稱；I：Inputs 鍵→名稱）。"""

    def __init__(self, write_registry=True):
        self.src = yaml.safe_load(SRC_YAML.read_text(encoding="utf-8"))["items"]
        self.inp = yaml.safe_load(INP_YAML.read_text(encoding="utf-8"))["items"]
        reg = json.loads(REGISTRY_PATH.read_text(encoding="utf-8")) if REGISTRY_PATH.exists() else {}
        sreg = reg.setdefault("SRC", {})
        for r in self.src:                         # SRC_ZP ID 由 Z1 發出：只核對、不改號
            if r["key"] in sreg and sreg[r["key"]] != r["id"]:
                raise SystemExit(f"SRC_ZP ID 與 registry 不符：{r['key']} {r['id']} ≠ {sreg[r['key']]}")
            if r["id"] in sreg.values() and sreg.get(r["key"]) != r["id"]:
                raise SystemExit(f"SRC_ZP ID 重用：{r['id']}")
            sreg[r["key"]] = r["id"]
        ireg = reg.setdefault("INP", {})
        retired = set(reg.setdefault("retired_INP", []))
        nxt = max([int(v[4:]) for v in list(ireg.values()) + list(retired)] + [0]) + 1
        keys = [r["key"] for r in self.inp]
        assert len(set(keys)) == len(keys), "Inputs 鍵重複"
        for r in self.inp:
            if r["key"] not in ireg:
                ireg[r["key"]] = f"INP_{nxt:03d}"
                nxt += 1
            r["id"] = ireg[r["key"]]
            assert r["tag"] in TAGS, (r["key"], r["tag"])
        for k in [k for k in ireg if k not in set(keys)]:      # 移除的鍵：號碼退役、不重用
            reg["retired_INP"].append(ireg.pop(k))
        self.reg = reg
        if write_registry:
            self.save_registry()
        self.src_by_key = {r["key"]: r for r in self.src}
        self.inp_by_key = {r["key"]: r for r in self.inp}
        self.inp_rowno = {r["key"]: 5 + i for i, r in enumerate(self.inp)}

    def save_registry(self):
        REGISTRY_PATH.write_text(json.dumps(self.reg, ensure_ascii=False, indent=0), encoding="utf-8")

    def S(self, key, part="value"):
        r = self.src_by_key[key]
        if not _num(r[part]):
            raise KeyError(f"SRC {key} 的 {part} 不是數值")
        return r["id"] + {"value": "", "lo": "_Lo", "hi": "_Hi"}[part]

    def I(self, key):
        return self.inp_by_key[key]["id"]

    def I_cell(self, key, col):
        """Inputs 的低（G）或高（H）欄位址。"""
        return f"Inputs!${col}${self.inp_rowno[key]}"


# ───────────────────────── 各頁 ─────────────────────────
def sheet_readme(wb, D, snap, oai, date, summary):
    ws = wb.active
    ws.title = "README"
    title(ws, f"智譜（智譜華章，02513.HK）收支模型 {VERSION}（{date}）",
          "命題（規格 D1，與 OpenAI v0.6 同一句）：智譜每 VR 等值 GW 的年營收能否覆蓋每 GW 年全成本；若不能，缺口要多少外部資金、由誰以什麼條件提供。FY2025–FY2030，曆年制。")
    lines = [
        ("本版範圍", "v0.1-Z3：SRC_ZP（公司財務原始數據）、TK_Link（Tokenomics 快照）、OAI_Link（OpenAI v0.6 命題輸出快照）、Inputs（假設）、Demand（需求與 token）、Revenue（營收）、"
                     "Compute（算力）、Cost（成本與命題表）、Checks。Funding、Reverse（Z4）尚未建立。"),
        ("幣別與單位", "公司金額一律人民幣億元（RMB 億），與財報一致。港幣、美元數據以 Inputs 匯率（USD/CNY、HKD/CNY，2026-09-28 中間價；定義常數，區間 ±5%）換算後才進計算頁。"
                     "API 牌價：人民幣元／百萬 token；token：兆（T）；Coding Plan 訂閱者：萬人；智譜清言用戶：百萬人。OAI_Link 為美元（$B），不換算、不參與計算。"),
        ("Excel 為唯一計算引擎", "藍字＝輸入（Excel 擁有）；黑字＝公式；綠字＝跨頁連結。builder 只產生結構，重建時保留 Excel 內已改過的藍字（隱藏頁 _Defaults）。公式不含常數（E6）：單位換算、天數、月數為 Inputs 的定義常數。"),
        ("SRC_ZP", f"{summary['src']} 列（SRC_ZP_001–{summary['src']:03d}，Z1 發出、穩定不改號）；有數值者具名 SRC_ZP_nnn，有低／高者另具名 _Lo／_Hi。來源四欄為 CC 初評（builder/source_rules.py）。"),
        ("TK_Link", f"Tokenomics {snap['version']}（{snap['file']}），master 提交 {snap['sha'][:7]}，讀取日 {date}；{summary['tk_ok']} 個具名範圍（同 OpenAI v0.6 的 63 名）＋{summary['tk_table']} 列以列標籤讀表（狀態『{TABLE_NOTE}』，只供對照）。"),
        ("OAI_Link", f"OpenAI v0.6 Excel（{oai['file']}，{oai['ref']}；SHA-256 {oai['sha'][:16]}…）的命題輸出逐年值（美元）；只被 Checks 與 HTML 引用（規格 D26）。"),
        ("Inputs", f"{summary['inp']} 列（INP_001 起，依 data/zhipu_inputs.yaml 的鍵發號）；每列附標記、依據、區間理由；Analogy／Assumed 一律給區間。"),
        ("Demand", "需求：①API 按量計費 token（2025、1H26＝營收 ÷ 有效單價倒推；2H26 起任務數 × 每任務 token × 價格反應）②GLM Coding Plan 訂閱者（方案別）× 每訂閱者 token ③智譜清言（免費）用戶 × 任務 × 每任務 token；"
                   "輸出 DEM_Tok_*（層級 Sol／Luna × 付費／免費）。本地化部署不耗智譜算力，不進 Demand（D4）。"),
        ("Revenue", "營收：①開放平台及 API（按量計費：層級 × 輸入／快取／輸出 × 有效價；Coding Plan 子列：方案別訂閱）②企業級智能體 ③企業級通用大模型 ④技術服務及其他（②③④＝本地化部署分部，規格 D2r）；"
                    "廣告＝0（D3）；總額＝淨額（D5）；雲端 vs 本地化、企業 vs 個人彙總；2025 與 1H26 對財報校準差距＝0；公司說法（ARR、平均售價 +101%、呼叫量 40 倍）只列對照。"),
        ("Compute", "晶片族（Hopper＝TK Hopper H100 欄；H20、國產＝Hopper × Inputs 比例，Tokenomics 缺口）× 層級（Sol／Luna）每 GW 年產能 → token 換算推論 GW；"
                    "2025、1H26 算力服務費（推論＝API 銷售成本、研發＝研發開支扣股權報酬 × 占比）÷ 每 GW 價格（卡時租價 × 每 GW 卡數 × 小時）→ 供給 GW 與 η；"
                    "2H26 起有效推論 GW＝token GW ÷ η，供給依需求配置；研發 GW＝殘差；容量上限係數回乘雲端營收；VR 等值 GW；10 萬國產晶片對照；R4 對帳；敏感度。"),
        ("Cost", "算力成本（2025、1H26＝算力服務費實付；之後＝租用 GW × 每 GW 價格＋自有資本支出，基準 0）、供應商持有成本與雲端毛利、本地化部署交付成本、"
                 "非算力成本（2025、1H26＝財報殘差；之後人數 × 每人成本）、股權報酬；命題表：每 VR 等值 GW 的營收、成本、差額、雲端口徑差額、覆蓋率（人民幣億與美元 $B）。"
                 "2025 與 1H26 全成本對財報費用合計差距＝0。"),
        ("Checks", "C01 起；CHK_Errors 必須為 0。WARN（ARR、η 合理範圍、國產晶片 GW 對照）與 INFO 不計入。"),
        ("標記", "Verified／Interested-party／Analogy／Assumed／Derived／Decision（Analogy、Assumed 一律附區間）。"),
        ("分層", "公司財務原始數據→SRC_ZP；AI 技術與算力→TK_Link（取自 Tokenomics）；假設→Inputs；OpenAI 對照→OAI_Link（不進計算）。"),
    ]
    for i, (a, b) in enumerate(lines):
        put(ws, f"A{4 + i}", a, F_BOLD)
        put(ws, f"B{4 + i}", b, F_CALC, wrap=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 130


def sheet_src(wb, D, final):
    ws = wb.create_sheet(SRC)
    title(ws, "SRC_ZP — 第 0 層 Source：智譜公司財務原始數據（財報、分部、牌價、用戶、算力採購、融資、股本與市值）",
          "只存已取得的原始訊息（非計算、非假設）。藍字＝原始值，由本活頁簿擁有；來源 data/zhipu_src.yaml（Z1）。AI 技術與算力的物理數據不在此登錄（取自 TK_Link）。",
          "來源等級（1 一手已讀原文／2 二手或轉載／3 待查核）、立場、立場說明、一手／二手：builder/source_rules.py 機械初評（欄 X 標『CC 初評』）。無數值列＝日期、事件或找不到數值的說明列（另見 Checks）。")
    header(ws, 4, SRC_COLS)
    for i, r in enumerate(D.src):
        n = 5 + i
        cl = source_rules.classify(r)
        r["_cl"] = cl
        vals = [r["id"], r["key"], r["metric"], None, None, None, r["unit"], r["period"], r["source"], r["url"], r["doc_date"], r["retrieved"],
                r["tag"], r["party"], r["basis"], cl["grade"], cl["stance"], cl["why"], cl["hand"], r["note"]]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=n, column=c, value=v if v is None or _num(v) else str(v)).font = F_CALC
        for col, part in zip("DEF", ("value", "lo", "hi")):
            v = final[(SRC, r["id"], col)]
            ws[f"{col}{n}"].value = v
            ws[f"{col}{n}"].font = F_IN
        put(ws, f"U{n}", f'=IF(AND($D{n}="",$E{n}="",$F{n}=""),1,0)')
        put(ws, f"V{n}", f'=IF(OR($I{n}="",$I{n}="—"),1,0)')
        put(ws, f"W{n}", f'=IF(AND(ISNUMBER($D{n}),ISNUMBER($E{n})),IF($E{n}>$D{n},1,0),0)+IF(AND(ISNUMBER($D{n}),ISNUMBER($F{n})),IF($D{n}>$F{n},1,0),0)'
                         f'+IF(AND(ISNUMBER($E{n}),ISNUMBER($F{n})),IF($E{n}>$F{n},1,0),0)')
        put(ws, f"X{n}", "CC 初評", F_NOTE)
        for col, part, suf in (("D", "value", ""), ("E", "lo", "_Lo"), ("F", "hi", "_Hi")):
            if _num(r[part]):
                nm(wb, r["id"] + suf, f"{SRC}!${col}${n}")
    for col, w in zip("ABCDEFGHI", (13, 26, 34, 10, 8, 8, 14, 14, 44)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D5"


def sheet_inputs(wb, D, final):
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs — 假設輸入（藍字）：值、低、高、標記、依據、區間理由；匯率與單位換算定義常數（D23、E6）",
          "長表格式：一列一個參數（× 索引）。來源 data/zhipu_inputs.yaml；INP_nnn 由 builder/id_registry.json 依鍵發出，穩定不重用。",
          "Analogy 寫明參照（如『OpenAI v0.6 INP_nnn』）與差異；Assumed 寫明依據；兩者一律給區間。Decision＝定義常數或規格決策（低＝值＝高）。")
    header(ws, 4, INP_COLS)
    for i, r in enumerate(D.inp):
        n = 5 + i
        vals = [r["id"], r["key"], r["name"], str(r["index"]), r["unit"], None, None, None, r["tag"], r["basis"], r["range"]]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=n, column=c, value=v).font = F_CALC
        for col in "FGH":
            ws[f"{col}{n}"].value = final[("Inputs", r["id"], col)]
            ws[f"{col}{n}"].font = F_IN
        put(ws, f"L{n}", f'=IF(AND(ISNUMBER($F{n}),ISNUMBER($G{n}),ISNUMBER($H{n})),IF(AND($G{n}<=$F{n},$F{n}<=$H{n}),0,1),0)')
        put(ws, f"M{n}", f'=IF(AND(OR($I{n}="Analogy",$I{n}="Assumed"),OR($G{n}="",$H{n}="")),1,0)')
        nm(wb, r["id"], f"Inputs!$F${n}")
    for col, w in zip("ABCDEFGHIJK", (10, 18, 46, 12, 12, 10, 9, 9, 12, 70, 36)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "F5"


def sheet_tk(wb, snap, date):
    ws = wb.create_sheet("TK_Link")
    title(ws, "TK_Link — Tokenomics 快照（第 0 層連結；不用 Excel 外部連結）",
          "具名區（同 OpenAI v0.6 的 63 名：IF_／SRC_DEM_／L1_）：值（藍字）由 builder 從 Tokenomics master 的 model/CURRENT 讀出寫入。本模型公式只引用 TK_ 名稱或本頁儲存格。",
          f"狀態『{TABLE_NOTE}』：Tokenomics 無具名範圍、以列標籤讀表（Cap_In F 表的 GLM 列、Price_Frontier 的 GLM 列與中國合格前緣）；只供 Revenue 對照列與 Checks，不作驅動（工作單 Z2 第 2 步）。")
    meta = [("Tokenomics 檔案", snap["file"], "TK_File"), ("Tokenomics 版本", snap["version"], "TK_Version"),
            ("Tokenomics 提交 SHA（master）", snap["sha"], "TK_Commit"), ("讀取日期", date, "TK_ReadDate")]
    for i, (lab, v, name) in enumerate(meta):
        put(ws, f"A{3 + i}", lab, F_BOLD)
        put(ws, f"B{3 + i}", v, F_CALC)
        nm(wb, name, f"TK_Link!$B${3 + i}")
    put(ws, "I7", "世代", F_BOLD)
    put(ws, "I8", "成本情境", F_BOLD)
    for k in range(15):
        put(ws, f"{L(TK_VCOL0 + k)}7", snap["hdr_gen"][k], F_CALC)
        put(ws, f"{L(TK_VCOL0 + k)}8", snap["hdr_cost"][k], F_CALC)
    nm(wb, "TK_HdrGen", "TK_Link!$J$7:$X$7")
    nm(wb, "TK_HdrCost", "TK_Link!$J$8:$X$8")
    header(ws, 9, ["Tokenomics 名稱", "本模型名稱", "說明", "單位", "欄數", "狀態", "Tokenomics 版本", "提交 SHA", "讀取日期"] + [f"值{k + 1}" for k in range(15)])
    r = TK_FIRST
    cells = {}
    for row in snap["rows"] + snap["table"]:
        n = len(row["values"])
        tkname = row["name"] if row["kind"] != "RD" else f"（讀表）{row.get('sheet')}!{row.get('row')}"
        vals = [tkname, "TK_" + row["name"], row["label"], row["unit"], n, row["status"], snap["version"], snap["sha"][:7], date]
        for c, v in enumerate(vals, start=1):
            put(ws, f"{L(c)}{r}", v, F_CALC)
        for k, v in enumerate(row["values"]):
            put(ws, f"{L(TK_VCOL0 + k)}{r}", v, F_IN)
        if n:
            nm(wb, "TK_" + row["name"], f"TK_Link!$J${r}" + (f":${L(TK_VCOL0 + n - 1)}${r}" if n > 1 else ""))
        cells[row["name"]] = [f"TK_Link!${L(TK_VCOL0 + k)}${r}" for k in range(n)]
        r += 1
    for row in snap["pending"]:
        vals = [row["name"], "（Tokenomics 現行版無此名稱；無具名範圍）", "Tokenomics 缺口：見報告", "", 0, row["status"], snap["version"], snap["sha"][:7], date]
        for c, v in enumerate(vals, start=1):
            put(ws, f"{L(c)}{r}", v, F_CALC)
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 60
    ws.freeze_panes = "D10"
    return cells


def sheet_oai(wb, oai, date):
    ws = wb.create_sheet("OAI_Link")
    title(ws, "OAI_Link — OpenAI v0.6 命題輸出快照（美元；只被 Checks 與 HTML 引用，規格 D26）",
          "值（藍字）由 builder 從 OpenAI v0.6 Excel 的具名範圍讀出（LibreOffice 重算後快取值）。本頁不參與任何計算；test_builder 掃描其他頁公式不得含 OAI_。",
          "比較問題：同樣以 token 計價、單價只有美國前沿模型幾分之一、算力用國產晶片的中國模型公司，每 GW 經濟與 OpenAI 差多少（Z5 HTML 並排）。")
    meta = [("OpenAI Excel 檔案", oai["file"], "OAI_File"), ("SHA-256", oai["sha"], "OAI_SHA256"), ("來源分支與提交", oai["ref"], "OAI_Ref"), ("讀取日期", date, "OAI_ReadDate")]
    for i, (lab, v, name) in enumerate(meta):
        put(ws, f"A{3 + i}", lab, F_BOLD)
        put(ws, f"B{3 + i}", v, F_CALC)
        nm(wb, name, f"OAI_Link!$B${3 + i}")
    header(ws, 9, ["OpenAI 名稱", "本模型名稱", "說明", "單位", "狀態", "2025", "2026", "2027", "2028", "2029", "2030"])
    for i, row in enumerate(oai["rows"]):
        r = OAI_FIRST + i
        for c, v in enumerate([row["name"], "OAI_" + row["name"], row["label"], row["unit"], row["status"]], start=1):
            put(ws, f"{L(c)}{r}", v, F_CALC)
        for k, v in enumerate(row["values"]):
            put(ws, f"{L(6 + k)}{r}", v, F_IN)
        if row["values"]:
            nm(wb, "OAI_" + row["name"], f"OAI_Link!$F${r}:$K${r}")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 34
    return len(oai["rows"])


def sheet_checks(wb, D, rows):
    """rows：(label, formula, expected, kind, note)；kind＝eq／tol／tolr（四捨五入容差）／warn／info。"""
    ws = wb.create_sheet("Checks")
    title(ws, f"Checks — {VERSION} 檢查（公式；結果 ERR 的格數＝CHK_Errors，必須為 0）",
          "OK／ERR：比對期望值；WARN：超出門檻只提示（不計入 CHK_Errors）；INFO：只列示。編號（C##）由 builder/id_registry.json 固定，新檢查取下一號，退役號碼不重用。期望值為 builder 寫入的常數。")
    header(ws, 4, ["編號", "檢查項", "值（公式）", "期望", "結果", "說明"])
    reg = D.reg
    known = reg.setdefault("CHK", {})
    retired = reg.setdefault("retired_CHK", [])
    nxt = max([int(v[1:]) for v in list(known.values()) + retired] + [0]) + 1
    ids = []
    for r_ in rows:
        if r_[0] not in known:
            known[r_[0]] = f"C{nxt:02d}"
            nxt += 1
        ids.append(known[r_[0]])
    assert len(set(ids)) == len(ids)
    labels = {r_[0] for r_ in rows}
    for k in [k for k in known if k not in labels]:
        retired.append(known.pop(k))
    order = sorted(range(len(rows)), key=lambda i: int(ids[i][1:]))
    out = []
    for j, i in enumerate(order):
        lab, f, exp, kind, note = rows[i]
        n = 5 + j
        put(ws, f"A{n}", ids[i], F_CALC)
        put(ws, f"B{n}", lab, F_CALC)
        put(ws, f"C{n}", f, F_CALC)
        if exp is not None and kind not in ("warnrng", "warnabs"):
            put(ws, f"D{n}", exp, F_CALC)
        if kind == "eq":
            put(ws, f"E{n}", f'=IF(C{n}=D{n},"OK","ERR")')
        elif kind == "tol":
            put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(ABS(C{n}-D{n})<0.000000001,"OK","ERR"),"ERR")')
        elif kind == "tolr":     # 公告四捨五入容差（億元第四位小數）
            put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(ABS(C{n}-D{n})<=0.0005,"OK","ERR"),"ERR")')
        elif kind == "warn":
            put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(ABS(C{n})>{D.I("arr_warn")},"WARN","OK"),"WARN")')
        elif kind == "warnrng":     # exp＝(下限 Inputs 鍵, 上限 Inputs 鍵)
            lo_, hi_ = D.I(exp[0]), D.I(exp[1])
            put(ws, f"D{n}", f"{lo_}～{hi_}", F_CALC)
            put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(AND(C{n}>={lo_},C{n}<={hi_}),"OK","WARN"),"WARN")')
        elif kind == "warnabs":     # exp＝門檻 Inputs 鍵
            put(ws, f"D{n}", f"={D.I(exp)}")
            put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(ABS(C{n})>{D.I(exp)},"WARN","OK"),"WARN")')
        else:
            put(ws, f"E{n}", "INFO")
        put(ws, f"F{n}", note, F_NOTE)
        out.append((ids[i], lab, kind))
    last = 4 + len(rows)
    put(ws, f"A{last + 2}", "CHK_Errors", F_BOLD)
    put(ws, f"C{last + 2}", f'=COUNTIF($E$5:$E${last},"ERR")')
    nm(wb, "CHK_Errors", f"Checks!$C${last + 2}")
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["F"].width = 70
    return ws, out


# ───────────────────────── 掃描（E6、OAI 隔離；Checks 與 test_builder 共用） ─────────────────────────
CALC_SHEETS = ("Demand", "Revenue", "Compute", "Cost", "Funding", "Reverse")


def e6_violations(ws):
    bad = []
    for row in ws.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                f = re.sub(r"'?[A-Za-z_]+'?!\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?", "", c.value)    # 跨頁參照
                f = re.sub(r"\$?\b[A-Z]{1,3}\$?\d+\b", "", f)                                          # 儲存格參照
                f = re.sub(r"\b[A-Za-z_][A-Za-z_0-9]*\b", "", f)                                          # 具名範圍與函數
                f = f.replace("(1-", "(").replace("+1)", ")").replace("(1+", "(")                        # 恆等式：1−比例、1＋成長率
                if re.search(r"\d", f):
                    bad.append(f"{ws.title}!{c.coordinate}: {c.value}")
    return bad


def oai_refs(wb):
    bad = []
    for ws in wb.worksheets:
        if ws.title in ("OAI_Link", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and ("OAI_" in c.value):
                    bad.append(f"{ws.title}!{c.coordinate}")
    return bad


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tk-dir", required=True, type=Path)
    ap.add_argument("--oai-xlsx", required=True, type=Path)
    ap.add_argument("--oai-ref", default="claude/openai-s6-release@e91df57")
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--base", type=Path)
    ap.add_argument("--date", default="2026-10-08")
    ap.add_argument("--no-lo", action="store_true", help="略過 LibreOffice 重算（僅供除錯；需在報告註明）")
    a = ap.parse_args()
    D = Data()
    snap = read_snapshot(a.tk_dir)
    oai = read_oai(a.oai_xlsx, a.oai_ref)
    outdir = ROOT / "build_out"
    outdir.mkdir(parents=True, exist_ok=True)
    a.out.parent.mkdir(parents=True, exist_ok=True)

    defaults = {}
    for r in D.src:
        for col, k in zip("DEF", ("value", "lo", "hi")):
            defaults[(SRC, r["id"], col)] = r[k]
    for r in D.inp:
        for col, k in zip("FGH", ("value", "lo", "hi")):
            defaults[("Inputs", r["id"], col)] = r[k]
    if a.base:
        final, matched, kept, unmatched = preserve.merge(defaults, a.base, outdir / "restore_log.txt")
    else:
        final, matched, kept, unmatched = dict(defaults), 0, 0, []
        (outdir / "restore_log.txt").write_text("matched: 0; Excel value kept over code default: 0; unmatched: 0\n（首次建置，無底稿）\n", encoding="utf-8")
    print(f"restore: matched {matched}, Excel kept over code {kept}, unmatched {len(unmatched)}")

    wb = openpyxl.Workbook()
    summary = dict(src=len(D.src), inp=len(D.inp), tk_ok=len(snap["rows"]), tk_table=sum(r["status"] == TABLE_NOTE for r in snap["table"]),
                   tk_pending=len(snap["pending"]), oai=len(oai["rows"]))
    sheet_readme(wb, D, snap, oai, a.date, summary)
    sheet_src(wb, D, final)
    tk_cells = sheet_tk(wb, snap, a.date)
    sheet_oai(wb, oai, a.date)
    sheet_inputs(wb, D, final)
    Z = z2.build(wb, D, tk_cells)
    Z.update(z3.build(wb, D, Z, snap))
    z2.fill(wb, Z)
    # 結構掃描（結果寫入 Checks 作為常數；test_builder 另行逐格驗證）
    e6 = sum(len(e6_violations(wb[s])) for s in CALC_SHEETS if s in wb.sheetnames)
    oref = len(oai_refs(wb))
    rows = z2.checks(D, Z, summary, snap, e6, oref) + z3.checks(D, Z)
    sheet_checks(wb, D, rows)
    z2.fill(wb, Z)
    D.save_registry()
    preserve.write_defaults(wb, final)
    raw = Path(tempfile.mkdtemp()) / a.out.name
    wb.save(raw)
    if a.no_lo:
        raw.replace(a.out)
    else:
        lo_recalc(raw, Path(tempfile.mkdtemp())).replace(a.out)
    json.dump(dict(summary, tk_file=snap["file"], tk_sha=snap["sha"], oai_file=oai["file"], oai_sha=oai["sha"], e6=e6, oai_refs=oref),
              open(outdir / "build_summary.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("saved", a.out, summary, "e6", e6, "oai_refs", oref)


if __name__ == "__main__":
    main()
