"""builder 結構性測試（不需 LibreOffice）：穩定 ID、yaml 與 Excel 一致、E6、OAI_Link 隔離、D4（本地化部署不進 Demand）。"""
import json
import re
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "builder"))
import build  # noqa: E402
from engine import Engine, current_model_path  # noqa: E402

EXPECT = yaml.safe_load((Path(__file__).parent / "scenarios.yaml").read_text(encoding="utf-8"))["workbook_expectations"]
SRC = yaml.safe_load((REPO / "data" / "zhipu_src.yaml").read_text(encoding="utf-8"))["items"]
INP = yaml.safe_load((REPO / "data" / "zhipu_inputs.yaml").read_text(encoding="utf-8"))["items"]
REG = json.loads((REPO / "builder" / "id_registry.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(current_model_path())


def _same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-12 * max(abs(a), abs(b), 1)
    return a == b


def test_src_ids_stable(wb):
    """SRC_ZP_nnn 由 Z1 發出：yaml、registry、Excel 三者一致，連號 001–604、不重用。"""
    assert [r["id"] for r in SRC] == [f"SRC_ZP_{i:03d}" for i in range(1, EXPECT["src_rows"] + 1)]
    assert REG["SRC"] == {r["key"]: r["id"] for r in SRC}
    ws = wb["SRC_ZP"]
    got = [(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(5, 5 + len(SRC))]
    assert got == [(r["id"], r["key"]) for r in SRC]


def test_inp_ids_stable(wb):
    """INP_nnn 依鍵發號：registry 與 Excel 一致；退役號碼不在使用中；號碼不重複。"""
    ids = list(REG["INP"].values())
    assert len(set(ids)) == len(ids)
    assert not set(ids) & set(REG.get("retired_INP", []))
    assert set(REG["INP"]) == {r["key"] for r in INP}
    ws = wb["Inputs"]
    got = {ws.cell(r, 2).value: ws.cell(r, 1).value for r in range(5, 5 + len(INP))}
    assert got == REG["INP"]


def test_chk_ids_stable(wb):
    ids = list(REG["CHK"].values())
    assert len(set(ids)) == len(ids)
    assert not set(ids) & set(REG.get("retired_CHK", []))
    ws = wb["Checks"]
    got = {ws.cell(r, 2).value: ws.cell(r, 1).value for r in range(5, 5 + len(ids))}
    assert got == REG["CHK"]


def test_src_and_inputs_values_equal_yaml():
    """有數值的列：Excel 具名範圍的值＝yaml（Excel 優先：若 Excel 改過，需先回寫 yaml 或在報告說明）。"""
    eng = Engine(current_model_path())
    for r in SRC:
        for part, suf in (("value", ""), ("lo", "_Lo"), ("hi", "_Hi")):
            if isinstance(r[part], (int, float)):
                assert _same(eng.get_name(r["id"] + suf), r[part]), (r["id"], part)
    for r in INP:
        assert _same(eng.get_name(REG["INP"][r["key"]]), r["value"]), r["key"]


def test_inputs_tags_and_ranges():
    """共同規則第 4 節：Analogy／Assumed 一律給區間，且低 ≤ 值 ≤ 高；每列有依據與區間理由。"""
    for r in INP:
        assert r["tag"] in build.TAGS, r["key"]
        assert r.get("basis") and r.get("range"), r["key"]
        if r["tag"] in ("Analogy", "Assumed"):
            assert isinstance(r["lo"], (int, float)) and isinstance(r["hi"], (int, float)), r["key"]
        if all(isinstance(r[k], (int, float)) for k in ("lo", "value", "hi")):
            assert r["lo"] <= r["value"] <= r["hi"], r["key"]


def test_e6_no_constants_in_calc_sheets(wb):
    """E6：Demand、Revenue 的公式不得內含常數（恆等式 1−比例、1＋成長率除外；單位換算、天數、月數在 Inputs）。"""
    for s in build.CALC_SHEETS:
        if s in wb.sheetnames:
            assert build.e6_violations(wb[s]) == [], s
    n = sum(1 for s in ("Demand", "Revenue") for row in wb[s].iter_rows(min_row=5) for c in row
            if isinstance(c.value, str) and c.value.startswith("="))
    assert n >= 600


def test_oai_link_isolated(wb):
    """規格 D26：OAI_Link 只被 Checks（與 HTML）引用；其他頁公式不得含 OAI_ 或 OAI_Link!。"""
    assert build.oai_refs(wb) == []
    for ws in wb.worksheets:
        if ws.title in ("OAI_Link", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "OAI_Link!" not in c.value, f"{ws.title}!{c.coordinate}"


def test_onprem_not_in_demand(wb):
    """規格 D4：本地化部署（②③④）不耗智譜算力，不進 Demand。Demand 公式不得引用其 SRC 列或 REV_OnPrem／REV_Other。"""
    keys = [r["id"] for r in SRC if re.match(r"rev_(agent|gpllm|techsvc|onprem)_", r["key"])]
    for row in wb["Demand"].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert not any(re.search(rf"\b{k}\b", c.value) for k in keys), c.coordinate
                assert "REV_OnPrem" not in c.value and "REV_Other" not in c.value, c.coordinate


def test_no_external_links_and_allowed_names(wb):
    """不使用 Excel 外部連結；計算頁引用的具名範圍只限 TK_／SRC_ZP_／INP_ 與本模型 DEM_／REV_ 名稱。"""
    names = set(wb.defined_names.keys())
    for s in ("Demand", "Revenue"):
        for row in wb[s].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "[" not in c.value, c.coordinate
                    for tok in re.findall(r"\b[A-Za-z_][A-Za-z0-9_]*\b", c.value):
                        if tok in names:
                            assert tok.startswith(("TK_", "SRC_ZP_", "INP_", "DEM_", "REV_")), (s, c.coordinate, tok)


def test_tk_table_rows_not_drivers(wb):
    """工作單 Z2 第 2 步：讀表列只供 Revenue 對照列與 Checks；Demand 不得引用 TK_Link。"""
    for row in wb["Demand"].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert "TK_Link!" not in c.value and "TK_RD_" not in c.value, c.coordinate
    rows = {r: wb["Revenue"].cell(r, 2).value for r in range(7, wb["Revenue"].max_row + 1)}
    for r, lab in rows.items():
        for col in "DEFGHIKL":
            v = wb["Revenue"][f"{col}{r}"].value
            if isinstance(v, str) and "TK_Link!" in v:
                assert str(lab).startswith("TK："), (r, lab)
