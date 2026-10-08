"""智譜收支模型公式引擎：以 pycel 直接載入 Excel 活頁簿並重算。

規則（CLAUDE.md 第 1、2 節）：
- 不手抄任何公式；所有數值都由活頁簿內的公式計算。
- 輸入與輸出優先以具名範圍存取；測試才使用儲存格位址。
- 遇到引擎不支援的函數時，回報給 Andy，不得在此另寫旁路計算。
"""
from __future__ import annotations

import hashlib
import json
import platform
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from pycel import ExcelCompiler

def _patch_pycel_defined_names() -> None:
    """相容性轉接（非計算）：pycel 1.0b30 讀取具名範圍用 openpyxl 3.0 的 `defined_names.definedName` 清單，
    openpyxl 3.1 起改為 dict，公式內一旦使用具名範圍（v5.8 Block 4 起）即 AttributeError。
    此處只改「怎麼讀出名稱的目的地」，語意與 pycel 原實作相同；不涉及任何數值計算。"""
    from pycel.excelwrapper import ExcelOpxWrapper

    def defined_names(self):
        if self.workbook is not None and self._defined_names is None:
            self._defined_names = {}
            dn = self.workbook.defined_names
            items = dn.values() if hasattr(dn, "values") else dn.definedName
            for d_name in items:
                destinations = [(alias, wksht) for wksht, alias in d_name.destinations if wksht in self.workbook]
                if destinations:
                    self._defined_names[str(d_name.name)] = destinations
        return self._defined_names

    ExcelOpxWrapper.defined_names = property(defined_names)


_patch_pycel_defined_names()

REPO_ROOT = Path(__file__).resolve().parent.parent
MODEL_DIR = REPO_ROOT / "model"
MODEL_NAME_RE = re.compile(r"^\d{8}_Zhipu_v\d+(\.\d+)?\.xlsx$")
_REF_RE = re.compile(
    r"^(?:'(?P<q>(?:[^']|'')+)'|(?P<u>[^!']+))!"
    r"\$?(?P<c1>[A-Z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]+)\$?(?P<r2>\d+))?$"
)


def current_model_path(model_dir: Path = MODEL_DIR) -> Path:
    """model/ 內唯一一份現行活頁簿（舊版在 model/archive/）。"""
    files = sorted(p for p in model_dir.glob("*.xlsx") if MODEL_NAME_RE.match(p.name))
    if len(files) != 1:
        raise RuntimeError(
            f"model/ 必須恰有一份現行 xlsx（YYYYMMDD_Zhipu_vN.xlsx），實際為 {[p.name for p in files]}"
        )
    return files[0]


def parse_ref(attr_text: str) -> tuple[str, str]:
    """'Interface!$C$19:$Q$19' → ('Interface', 'C19:Q19')；頁名可含中文或引號。"""
    m = _REF_RE.match(attr_text)
    if not m:
        raise ValueError(f"無法解析範圍：{attr_text!r}")
    sheet = (m["q"] or m["u"]).replace("''", "'")
    ref = f"{m['c1']}{m['r1']}"
    if m["c2"]:
        ref += f":{m['c2']}{m['r2']}"
    return sheet, ref


_KEY_RE = re.compile(r"^(?P<name>[A-Za-z_][A-Za-z0-9_.]*)(?:\[(?P<i>\d+)\])?$")


def resolve_key(names: dict[str, str], key: str) -> tuple[str, str]:
    """情境輸入鍵 → (頁名, 儲存格)。鍵可為 'Sheet!A1'（舊情境）、'NAME'（單格具名範圍）
    或 'NAME[k]'（單列或單欄具名範圍的第 k 格，k 從 1 起算）。"""
    if "!" in key:
        sheet, coord = key.split("!")
        return sheet, coord
    m = _KEY_RE.match(key)
    if not m or m["name"] not in names:
        raise KeyError(f"找不到具名範圍：{key!r}")
    sheet, ref = parse_ref(names[m["name"]])
    if ":" not in ref:
        if m["i"]:
            raise ValueError(f"{key}：單格具名範圍不可加索引")
        return sheet, ref
    if not m["i"]:
        raise ValueError(f"{key}：多格具名範圍必須指定索引，例如 {m['name']}[1]")
    c1, c2 = ref.split(":")
    (a, r1), (b, r2) = (re.match(r"([A-Z]+)(\d+)", x).groups() for x in (c1, c2))
    k = int(m["i"]) - 1
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
    if r1 == r2:   # 單列：沿欄展開
        return sheet, f"{gl(ci(a) + k)}{r1}"
    if a == b:     # 單欄：沿列展開
        return sheet, f"{a}{int(r1) + k}"
    raise ValueError(f"{key}：只支援單列或單欄具名範圍")


def read_defined_names_xml(path: Path) -> dict[str, str]:
    """直接解析 workbook.xml 的 definedName，作為獨立於 openpyxl 的對照來源。"""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("xl/workbook.xml"))
    out = {}
    for el in root.iterfind(".//m:definedNames/m:definedName", ns):
        if el.get("localSheetId") is None:  # 只取活頁簿層級名稱
            out[el.get("name")] = (el.text or "").strip()
    return out


def norm(v):
    """統一空值：None 與空字串視為同一種空值；錯誤值 pycel 以 '#XXX!' 字串回傳。
    v5.11：SUMPRODUCT 回傳 numpy 純量（np.int64），轉為 Python 型別（只改型別表示，不改數值）。"""
    if v is None:
        return ""
    if hasattr(v, "item") and getattr(v, "shape", None) == ():
        return v.item()
    if isinstance(v, float) and type(v) is not float:        # 從序列化快取載入時為 ruamel 的 ScalarFloat／ScalarInt
        return float(v)
    if isinstance(v, int) and not isinstance(v, bool) and type(v) is not int:
        return int(v)
    return v


PLUGINS = ["engine.excel_semantics"]


def _cache_fingerprint(path: Path) -> dict:
    """序列化快取的有效條件：活頁簿內容、引擎原始碼、pycel 與 Python 版本任一改變即失效。"""
    import pycel
    here = Path(__file__).resolve().parent
    h = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()   # noqa: E731
    return {"xlsx": h(path), "core": h(here / "core.py"), "semantics": h(here / "excel_semantics.py"),
            "pycel": getattr(pycel, "__version__", "?"),
            # 只比到 major.minor：CI 各 runner 的 setup-python 可能給不同 patch 版（2026-10-07 實例 3.11.16／3.11.17），pickle 格式不受 patch 影響
            "python": ".".join(platform.python_version_tuple()[:2])}


class Engine:
    """封裝 pycel：載入、設定輸入、依具名範圍取值。

    cache：pycel 序列化檔（見 `save_cache`）。有給時從檔案載入計算圖，省去約 50 秒建圖；
    快取與活頁簿／引擎原始碼／pycel／Python 版本不符時直接報錯（不靜默改走重建）。"""

    def __init__(self, path: str | Path | None = None, cache: str | Path | None = None):
        self.path = Path(path) if path else current_model_path()
        if cache:
            self._xl = self._load_cache(Path(cache), self.path)
        else:
            self._xl = ExcelCompiler(filename=str(self.path), plugins=PLUGINS)
        wb = openpyxl.load_workbook(self.path)  # 公式模式：取得公式格清單與具名範圍
        self.sheetnames = list(wb.sheetnames)
        self.names = {k: v.attr_text for k, v in wb.defined_names.items()}
        self._formula_cells = [
            (ws.title, c.coordinate)
            for ws in wb
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        self._warm_up()     # 從快取載入時圖已建好，此處只剩強制全簿重算

    @staticmethod
    def _load_cache(cache: Path, path: Path):
        want = json.loads(cache.with_suffix(cache.suffix + ".json").read_text(encoding="utf-8"))
        have = _cache_fingerprint(path)
        if want != have:
            raise RuntimeError(f"引擎快取已失效，請重建：{ {k: (want.get(k), have[k]) for k in have if want.get(k) != have[k]} }")
        sys.setrecursionlimit(max(sys.getrecursionlimit(), 100000))
        return ExcelCompiler.from_file(str(cache), plugins=PLUGINS)

    def save_cache(self, cache: str | Path) -> None:
        """建圖並重算後，把 pycel 計算圖序列化到檔案，附 `.json` 指紋。"""
        cache = Path(cache)
        for name in self.names:      # 只靠具名範圍讀取的常數格（不被任何公式引用）也要先進計算圖，否則載入後讀到空值
            self._xl.evaluate(self._addr(*self.name_ref(name)))
        sys.setrecursionlimit(max(sys.getrecursionlimit(), 100000))
        self._xl.to_file(str(cache))
        cache.with_suffix(cache.suffix + ".json").write_text(json.dumps(_cache_fingerprint(self.path)), encoding="utf-8")

    def _warm_up(self) -> None:
        """建立全簿計算圖，並強制全部公式格重算。

        pycel 冷啟動時，公式格會直接回傳檔內的快取值而不計算（實測：recalculate() 後
        3,729 格的浮點位元改變）。若不強制重算，parity 的基準情境只是在比對快取值。
        本引擎因此一律在建構時重算，之後所有取值都是引擎自行計算的結果。
        """
        addrs = [self._addr(s, c) for s, c in self._formula_cells]
        self._xl.evaluate(addrs)      # 建圖；同時滿足 set_value 要求該格已在 cell map
        self._xl.recalculate()        # 清除快取值並重算

    # ── 基本存取 ────────────────────────────────────────────────
    @staticmethod
    def _addr(sheet: str, ref: str) -> str:
        return f"'{sheet}'!{ref}"

    def get(self, sheet: str, ref: str):
        """單格回傳純量；多格回傳巢狀 list（列 × 欄）。"""
        v = self._xl.evaluate(self._addr(sheet, ref))
        if isinstance(v, tuple):
            if v and not isinstance(v[0], tuple):  # pycel 對單列範圍回傳一維 tuple
                v = (v,)
            return [[norm(x) for x in row] for row in v]
        return norm(v)

    def set_input(self, sheet: str, coord: str, value) -> None:
        """改寫輸入格；相依格於下次取值時重算。"""
        addr = self._addr(sheet, coord)
        # pycel 限制：依賴格尚未建圖就設值，後建圖時會讀回檔內原值（F35 實測重現）；
        # 建構時已建全簿計算圖，故此處可直接設值。
        self._xl.set_value(addr, value)

    # ── 具名範圍 ────────────────────────────────────────────────
    def name_ref(self, name: str) -> tuple[str, str]:
        return parse_ref(self.names[name])

    def get_name(self, name: str):
        """單格→純量；單列或單欄→一維 list；其他→巢狀 list。"""
        sheet, ref = self.name_ref(name)
        v = self.get(sheet, ref)
        if isinstance(v, list) and (len(v) == 1 or all(len(r) == 1 for r in v)):
            return [x for row in v for x in row]
        return v

    def set_key(self, key: str, value) -> None:
        """依情境輸入鍵（位址、具名範圍或 NAME[k]）改寫輸入格。"""
        self.set_input(*resolve_key(self.names, key), value)

    def set_name(self, name: str, value) -> None:
        """只允許單格具名範圍（例：CTL_GW）。"""
        sheet, ref = self.name_ref(name)
        if ":" in ref:
            raise ValueError(f"{name} 不是單格具名範圍")
        self.set_input(sheet, ref, value)

    # ── 全簿 ────────────────────────────────────────────────────
    @property
    def formula_cells(self) -> list[tuple[str, str]]:
        return list(self._formula_cells)

    def evaluate_all(self) -> dict[tuple[str, str], object]:
        """重算並回傳全部公式格的值。"""
        addrs = [self._addr(s, c) for s, c in self._formula_cells]
        vals = self._xl.evaluate(addrs)
        return {k: norm(v) for k, v in zip(self._formula_cells, vals)}


__all__ = ["Engine", "resolve_key", "current_model_path", "parse_ref", "read_defined_names_xml", "norm"]
