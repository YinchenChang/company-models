#!/usr/bin/env python3
"""產生每段報告的對照 Excel（CLAUDE.md 第 5 節）：①本段新增或變動的每一列 ②本段關鍵輸出 FY2025–2030 ③已套用的預設。

用法：python3 tools/make_report_tables.py --defaults docs/reports/<報告>.md --out docs/reports/<報告>_對照.xlsx
值一律讀現行模型（LibreOffice 重算後的快取值；不另算）。③由報告 md 的「已套用的預設」表解析，報告與 Excel 同一來源。
"""
import argparse
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
from engine import current_model_path  # noqa: E402

BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="FFEDEDED")
WRAP = Alignment(wrap_text=True, vertical="top")
KEY_NAMES = [
    ("REV_API", "①開放平台及 API（按量計費＋Coding Plan）", "RMB 億"), ("REV_API_PayGo", "　其中按量計費", "RMB 億"),
    ("REV_Sub", "　其中 Coding Plan 子列（訂閱）", "RMB 億"), ("REV_Sub_Lite", "　　Lite", "RMB 億"), ("REV_Sub_Pro", "　　Pro", "RMB 億"),
    ("REV_Sub_Max", "　　Max", "RMB 億"), ("REV_OnPrem_Agent", "②企業級智能體", "RMB 億"), ("REV_OnPrem_GPLLM", "③企業級通用大模型", "RMB 億"),
    ("REV_Other", "④技術服務及其他", "RMB 億"), ("REV_OnPrem", "本地化部署（②＋③＋④）", "RMB 億"), ("REV_Ads", "廣告", "RMB 億"),
    ("REV_Gross", "營收總額（＝淨額）", "RMB 億"), ("REV_Cloud", "雲端營收（①）", "RMB 億"), ("REV_GrossCapped", "營收總額（截頂後；係數佔位 1）", "RMB 億"),
    ("REV_Enterprise", "企業", "RMB 億"), ("REV_Individual", "個人", "RMB 億"), ("REV_ApiPrice", "組合有效單價", "元／百萬 token"),
    ("DEM_Subs_CP", "Coding Plan 訂閱者（期間平均）", "萬人"), ("DEM_Users_Consumer", "智譜清言月活", "百萬人"),
    ("DEM_Tok_API", "API 按量計費 token", "T"), ("DEM_Tok_CP", "Coding Plan token", "T"), ("DEM_Tok_Consumer", "智譜清言 token", "T"),
    ("DEM_Tok_APIFree", "免費 API 型號 token", "T"), ("DEM_Tok_Paid_Sol", "付費 token：Sol", "T"), ("DEM_Tok_Paid_Luna", "付費 token：Luna", "T"),
    ("DEM_Tok_Free_Sol", "免費 token：Sol", "T"), ("DEM_Tok_Free_Luna", "免費 token：Luna", "T"), ("DEM_Tok_Total", "token 合計", "T"),
]
KEY_NAMES_Z3 = [
    ("CMP_SupplyGW", "供給 GW（租用＋自有）", "GW"), ("CMP_InfGW_Eff", "有效推論 GW（未截頂）", "GW"), ("CMP_InfGW", "推論 GW（截頂後）", "GW"),
    ("CMP_RDGW", "研發 GW（殘差）", "GW"), ("CMP_DemandGW", "總需求 GW", "GW"), ("CMP_TokGW", "推論 GW（token 換算）", "GW"),
    ("CMP_Eta", "η（逐年）", "倍"), ("CMP_RDShare", "研發占非閒置供給比例", "比例"), ("CMP_CapFactor", "容量上限係數", "倍"),
    ("CMP_VReqFactor", "機隊 VR 等值係數", "倍"), ("CMP_Supply_VReq", "供給 VR 等值 GW（命題分母）", "GW"),
    ("CMP_Share_Hopper", "晶片族占比：Hopper", "比例"), ("CMP_Share_H20", "晶片族占比：H20", "比例"), ("CMP_Share_Domestic", "晶片族占比：國產", "比例"),
    ("CMP_PricePerGW", "組合後每 GW 年價格", "RMB 億／GW／年"), ("CMP_DomesticGW", "國產供給 GW", "GW"),
    ("COST_Compute", "算力成本（現金口徑）", "RMB 億"), ("COST_InfCompute", "　其中推論", "RMB 億"), ("COST_RDCompute", "　其中研發", "RMB 億"),
    ("COST_SupplierHold", "供應商持有成本", "RMB 億"), ("COST_CloudGM", "雲端毛利", "RMB 億"), ("COST_CloudGMPct", "雲端毛利率", "比例"),
    ("COST_OnPremDelivery", "本地化部署交付成本", "RMB 億"), ("COST_Headcount", "員工人數（期間平均）", "人"),
    ("COST_NonCompExSBC", "非算力成本（不含股權報酬）", "RMB 億"), ("COST_SBC", "股權報酬", "RMB 億"), ("COST_FullCash", "全成本（含股權報酬）", "RMB 億"),
    ("REV_NetCapped", "營收淨額（截頂後）", "RMB 億"), ("COST_GapCash", "差額＝營收淨額 − 全成本", "RMB 億"),
]
PROP_RMB = [("COST_PropRev_VR", "營收淨額"), ("COST_PropCloudRev_VR", "雲端營收"), ("COST_PropCompute_VR", "算力成本"), ("COST_PropOnPrem_VR", "本地化部署交付成本"),
            ("COST_PropNonComp_VR", "非算力成本（不含股權報酬）"), ("COST_PropSBC_VR", "股權報酬"), ("COST_PropFull_VR", "全成本（含股權報酬）"),
            ("COST_PropFullExSBC_VR", "全成本（不含股權報酬）"), ("COST_PropGap_VR", "差額（含股權報酬）"), ("COST_PropGapExSBC_VR", "差額（不含股權報酬）"),
            ("COST_PropGap_VR_Cloud", "雲端口徑差額"), ("COST_Coverage", "覆蓋率（倍）"), ("COST_CoverageCloud", "雲端覆蓋率（倍）")]
PROP_USD = [("COST_PropRev_VR_USD", "營收淨額"), ("COST_PropCompute_VR_USD", "算力成本"), ("COST_PropFull_VR_USD", "全成本（含股權報酬）"),
            ("COST_PropGap_VR_USD", "差額（含股權報酬）"), ("COST_PropGap_VR_Cloud_USD", "雲端口徑差額"), ("COST_GapCash_USD", "差額（$B）")]
HALF_NAMES = [("REV_API_H", "①開放平台及 API"), ("REV_Sub_H", "Coding Plan 子列"), ("REV_OnPrem_H", "本地化部署"), ("REV_Gross_H", "營收總額"),
              ("DEM_Tok_API_H", "API 按量計費 token（T）"), ("REV_ApiPrice_H", "組合有效單價（元／M）")]


def _vals(wb, attr):
    m = re.match(r"^(?:'([^']+)'|([^!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", attr)
    ws = wb[m.group(1) or m.group(2)]
    from openpyxl.utils import column_index_from_string as ci
    c1, r1 = ci(m.group(3)), int(m.group(4))
    c2 = ci(m.group(5)) if m.group(5) else c1
    return [ws.cell(r1, c).value for c in range(c1, c2 + 1)]


def head(ws, labels):
    ws.append(labels)
    for c in ws[ws.max_row]:
        c.font, c.fill, c.alignment = BOLD, HDR, WRAP


def parse_defaults(md: Path):
    rows, on = [], False
    for line in md.read_text(encoding="utf-8").splitlines():
        if line.startswith("## ") and "已套用的預設" in line:
            on = True
            continue
        if on and line.startswith("## "):
            break
        if on and line.startswith("|") and not line.startswith("|---"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            rows.append(cells)
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--defaults", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--seg", default="z2", choices=("z2", "z3"))
    a = ap.parse_args()
    z3 = a.seg == "z3"
    sheets, codes = (("Compute", "Cost"), r"^[GK]\d+$") if z3 else (("Demand", "Revenue"), r"^[DR]\d+$")
    model = current_model_path()
    wbv = openpyxl.load_workbook(model, data_only=True)
    names = {k: v.attr_text for k, v in wbv.defined_names.items()}
    out = openpyxl.Workbook()

    # ① 本段新增的每一列
    ws = out.active
    ws.title = "①新增列_" + "_".join(sheets)
    head(ws, ["頁", "編號／ID", "列名", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "1H2026", "2H2026", "具名範圍", "來源／標記／說明"])
    for sh in sheets:
        s = wbv[sh]
        nm_by_row = {}
        for n, t in names.items():
            mm = re.match(rf"^{sh}!\$[A-Z]+\$(\d+)", t)
            if mm:
                nm_by_row.setdefault(int(mm.group(1)), []).append(n)
        for r in range(7, s.max_row + 1):
            code = s.cell(r, 1).value
            if not code or not re.match(codes, str(code)):
                continue
            vals = [s[f"{c}{r}"].value for c in "DEFGHIKL"]
            ws.append([sh, code, s.cell(r, 2).value, s.cell(r, 3).value] + vals + ["、".join(sorted(nm_by_row.get(r, []))), s[f"M{r}"].value])
    for col, w in zip("ABCDEFGHIJKLMN", (9, 18, 50, 12, 10, 10, 10, 10, 10, 10, 10, 10, 30, 70)):
        ws.column_dimensions[col].width = w
    ws = out.create_sheet("①新增列_Inputs")
    head(ws, ["INP_ID", "鍵", "參數", "索引", "單位", "值", "低", "高", "標記", "依據", "區間理由"])
    s = wbv["Inputs"]
    for r in range(5, s.max_row + 1):
        if s.cell(r, 1).value and (not z3 or int(str(s.cell(r, 1).value)[4:]) >= 74 or s.cell(r, 2).value in ("g_task_2h26", "onprem_k_2h26")):
            ws.append([s.cell(r, c).value for c in range(1, 12)])
    for col, w in zip("ABCDEFGHIJK", (9, 16, 46, 10, 12, 9, 9, 9, 11, 80, 36)):
        ws.column_dimensions[col].width = w
    ws = out.create_sheet("①新增列_TK_OAI_Checks")
    head(ws, ["頁", "名稱／編號", "說明", "單位／結果", "值1", "值2", "值3", "值4", "值5", "值6", "來源／期望／說明"])
    s = wbv["TK_Link"]
    for r in range(10, s.max_row + 1):
        if s.cell(r, 6).value == "讀表（非具名）" and not z3:
            ws.append(["TK_Link", s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value] + [s.cell(r, 10 + k).value for k in range(5)]
                      + [None, f"Tokenomics {s.cell(r, 7).value}（{s.cell(r, 8).value}）讀表，只供對照"])
    s = wbv["OAI_Link"]
    for r in range(10, s.max_row + 1):
        if s.cell(r, 1).value and not z3:
            ws.append(["OAI_Link", s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value] + [s.cell(r, 6 + k).value for k in range(6)]
                      + [f"OpenAI v0.6（{wbv['OAI_Link']['B5'].value}；SHA-256 {str(wbv['OAI_Link']['B4'].value)[:12]}…）；值為 2025–2030"])
    s = wbv["Checks"]
    for r in range(5, s.max_row + 1):
        if s.cell(r, 1).value and re.match(r"^C\d+$", str(s.cell(r, 1).value)) and (not z3 or int(str(s.cell(r, 1).value)[1:]) >= 43):
            ws.append(["Checks", s.cell(r, 1).value, s.cell(r, 2).value, s.cell(r, 5).value, s.cell(r, 3).value, None, None, None, None, None,
                       f"期望 {s.cell(r, 4).value}；{s.cell(r, 6).value}"])
    for col, w in zip("ABCDEFGHIJK", (9, 24, 60, 12, 10, 10, 10, 10, 10, 10, 70)):
        ws.column_dimensions[col].width = w

    # ② 關鍵輸出
    ws = out.create_sheet("②關鍵輸出")
    if z3:
        head(ws, ["命題表（人民幣；每 VR 等值 GW）", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "具名範圍"])
        ws.append(["分母：供給 VR 等值 GW", "GW"] + _vals(wbv, names["CMP_Supply_VReq"]) + ["CMP_Supply_VReq"])
        for n, lab in PROP_RMB:
            ws.append([lab, "倍" if n.startswith("COST_Coverage") else "RMB 億／GW／年"] + _vals(wbv, names[n]) + [n])
        ws.append([])
        head(ws, ["命題表（美元；每 VR 等值 GW；÷ USD/CNY ÷ 10）", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "具名範圍"])
        for n, lab in PROP_USD:
            ws.append([lab, "$B" if n == "COST_GapCash_USD" else "$B／GW／年"] + _vals(wbv, names[n]) + [n])
        ws.append(["對照：OpenAI v0.6 每 VR 等值 GW 差額", "$B／GW／年"] + _vals(wbv, names["OAI_COST_PropGap_VR"]) + ["OAI_COST_PropGap_VR"])
        ws.append([])
        head(ws, ["算力與成本", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "具名範圍"])
    for n, lab, unit in (KEY_NAMES_Z3 if z3 else KEY_NAMES):
        ws.append([lab, unit] + _vals(wbv, names[n]) + [n])
    ws.append([])
    if not z3:
        head(ws, ["2026 半年拆分", "", "1H2026（實際）", "2H2026（驅動）", "", "", "", "", "具名範圍"])
        for n, lab in HALF_NAMES:
            ws.append([lab, ""] + _vals(wbv, names[n]) + ["", "", "", "", n])
        ws.append([])
    head(ws, ["對照（只列差距，不反推）", "單位", "值", "", "", "", "", "", "具名範圍"])
    extra = (("CMP_Eta2025", "η 2025", "倍"), ("CMP_Eta1H26", "η 1H26", "倍"), ("CMP_ChipGW", "10 萬國產晶片換算 IT GW", "GW"),
             ("CMP_ChipGWGap", "模型 2H26 國產供給 GW ÷ 換算 GW − 1", "比例"), ("CMP_SpendGW", "2025 支出換算推論 GW", "GW"),
             ("CMP_R4Mult", "2025 Coding Plan token GW ÷ 支出換算推論 GW", "倍"), ("COST_Recon", "2025 全成本 − 財報費用合計", "RMB 億")) if z3 else ()
    for n, lab, unit in extra + (("REV_ARRGapMaaS", "模型 2H26 年化雲端營收 ÷ MaaS ARR − 1", "比例"), ("REV_ARRGapTotal", "模型 2H26 年化總營收 ÷ 全業務 ARR − 1", "比例"),
                         ("REV_Gap2025", "2025 營收總額校準差距", "RMB 億"), ("REV_Gap1H26", "1H26 營收總額校準差距（公告四捨五入）", "RMB 億"),
                         ("REV_GapUncal2025", "2025 未校準差距", "RMB 億"), ("DEM_ApiBilledRatio2025", "2025 計費比例", "比例"),
                         ("DEM_ImpliedVolGrowth", "2H25→1H26 實際隱含量成長", "倍"), ("CHK_Errors", "CHK_Errors", "格")):
        ws.append([lab, unit, _vals(wbv, names[n])[0], "", "", "", "", "", n])
    ws.column_dimensions["A"].width = 44
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 12

    # ③ 已套用的預設
    ws = out.create_sheet("③已套用的預設")
    rows = parse_defaults(a.defaults)
    for i, r in enumerate(rows):
        if i == 0:
            head(ws, r)
        else:
            ws.append(r)
    for col, w in zip("ABCDE", (6, 30, 60, 40, 70)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = WRAP
    a.out.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.out)
    print("saved", a.out, "defaults rows", len(rows) - 1)


if __name__ == "__main__":
    main()
