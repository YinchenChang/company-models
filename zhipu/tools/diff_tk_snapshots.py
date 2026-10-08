#!/usr/bin/env python3
"""兩份模型檔的 TK_Link 快照逐名比對（舊快照 vs 新快照）：列出值有變動者的逐欄前後值與變動率。

沿用舊 repo PR #2（提交 479e94b）的工具，S1 加上：逐欄變動率、表頭（世代／成本情境）標示、--json 輸出（供報告與對照 Excel）。
用法：python3 tools/diff_tk_snapshots.py <舊.xlsx> <新.xlsx> [--md out.md] [--json out.json]
"""
import argparse
import json
from pathlib import Path

import openpyxl

FIRST = 10          # TK_Link 表格起始列
VCOL0 = 10          # 值從 J 欄起


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True)["TK_Link"]
    hdr = [f"{ws.cell(7, VCOL0 + i).value}／{ws.cell(8, VCOL0 + i).value}" for i in range(15)]
    meta = {k: ws[f"B{r}"].value for k, r in (("file", 3), ("version", 4), ("sha", 5), ("date", 6))}
    out = {}
    for r in range(FIRST, ws.max_row + 1):
        n = ws.cell(r, 1).value
        if not n:
            continue
        k = ws.cell(r, 5).value or 0
        out[n] = dict(row=r, status=ws.cell(r, 6).value, label=ws.cell(r, 3).value, unit=ws.cell(r, 4).value,
                      values=[ws.cell(r, VCOL0 + i).value for i in range(k)])
    return meta, hdr, out


def same(p, q):
    if p == q:
        return True
    return isinstance(p, (int, float)) and isinstance(q, (int, float)) and abs(p - q) <= 1e-9 * max(abs(p), abs(q), 1e-300)


def rel(p, q):
    return (q - p) / p if isinstance(p, (int, float)) and isinstance(q, (int, float)) and p else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("old", type=Path)
    ap.add_argument("new", type=Path)
    ap.add_argument("--md", type=Path)
    ap.add_argument("--json", type=Path)
    a = ap.parse_args()
    mo, hdr_o, o = load(a.old)
    mn, hdr_n, n = load(a.new)
    changed, unchanged, removed = [], [], []
    for name, d in o.items():
        if d["status"] != "OK":
            continue
        if name not in n or n[name]["status"] != "OK":
            removed.append(name)
            continue
        nv = n[name]["values"]
        if len(nv) != len(d["values"]):
            changed.append(dict(name=name, unit=n[name]["unit"], label=n[name]["label"], note=f"欄數 {len(d['values'])}→{len(nv)}", cols=[]))
            continue
        cols = [dict(col=i + 1, hdr=(hdr_n[i] if len(nv) == 15 else ""), old=p, new=q, rel=rel(p, q))
                for i, (p, q) in enumerate(zip(d["values"], nv)) if not same(p, q)]
        (changed if cols else unchanged).append(dict(name=name, unit=n[name]["unit"], label=n[name]["label"], cols=cols, n=len(nv)))
    added = [dict(name=k, unit=v["unit"], label=v["label"], values=v["values"]) for k, v in n.items()
             if v["status"] == "OK" and (k not in o or o[k]["status"] != "OK")]
    res = dict(old=mo, new=mn, hdr_changed=hdr_o != hdr_n, hdr_old=hdr_o, hdr_new=hdr_n, changed=changed,
               unchanged=[u["name"] for u in unchanged], removed=removed, added=added)
    head = (f"舊快照 {mo['version']}（{mo['sha'][:7]}）→ 新快照 {mn['version']}（{mn['sha'][:7]}）："
            f"舊 OK {len(changed) + len(unchanged) + len(removed)} 名；值有變動 {len(changed)} 名；不變 {len(unchanged)} 名；"
            f"新快照已無 {len(removed)} 名；新增有值 {len(added)} 名；表頭變動：{'是' if res['hdr_changed'] else '否'}")
    lines = [head, "", "| 名稱 | 欄 | 世代／成本情境 | 舊值 | 新值 | 變動率 |", "|---|---|---|---|---|---|"]
    for c in changed:
        if not c["cols"]:
            lines.append(f"| {c['name']} | — | {c.get('note', '')} | | | |")
        for x in c["cols"]:
            lines.append(f"| {c['name']} | {x['col']} | {x['hdr']} | {x['old']:.6g} | {x['new']:.6g} | "
                         + (f"{x['rel']:+.2%}" if x["rel"] is not None else "—") + " |")
    print("\n".join(lines))
    if a.md:
        a.md.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if a.json:
        a.json.write_text(json.dumps(res, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
