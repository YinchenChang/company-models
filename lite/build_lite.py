# Lumentum（LITE，美股）收支與估值模型 v0.2.2（AI 半導體 Project：網通次產業內複製——雷射與光元件）
# 用法：python3 lite/build_lite.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度：6 月底附近的週六結束（FY2026＝2025-06-29～2026-06-27）。FY(t) 的後半年（1–6 月）在日曆年 t、前半年（7–12 月）在日曆年 t−1。
#   市場層（markets/network）是日曆年；本模型先在日曆年算各產品營收，再換成會計年度：
#   FY2027＝CY2026 下半年（FY2027 第一季指引＋第二季假設）＋ h × CY2027；FY2028＝（1 − h）× CY2027 ＋ h × CY2028；FY2029 同理（h＝日曆年上半年占比，Assumed）。
# 金額：十億美元（US$B；Q 頁為百萬美元）；EPS：美元。口徑：公司 non-GAAP（排除股票報酬、併購攤銷、可轉債清償損失等）；GAAP 只在 Q 頁註記。
# FY2026 錨點＝FY2026 全年實績；CY2026 錨點＝FY2026 第三、四季實績＋FY2027 第一季指引中點＋第二季假設（FY2027 第一季財報約 2026-11-05 公布後更新）。
import sys, os, glob, re
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
NWB = openpyxl.load_workbook(NET_F, data_only=True)
NETR = {}
for r in NWB["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in ("  其中 CPO", "  其中可插拔", "AI IT 資本支出（加速器 × 每顆 IT capex）"):
        NETR[lab.strip()] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
for r in NWB["Units"].iter_rows(min_row=4):
    if r[0].value == "Google TPU":
        NETR["TPU"] = ([r[1 + j].value for j in range(4)], f"{rel(NET_F)}｜Units!B{r[0].row}:E{r[0].row}")
assert len(NETR) == 4 and all(v is not None for k in NETR for v in NETR[k][0]), NETR.keys()
# 網通 CPO 三情境（markets/network v0.2.3；Andy 2026-10-25 13:13：技術替代屬行業情境，A＝替代快、C＝慢，公司讀同字母）
NETS = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    for s_ in ("A", "B", "C"):
        for k_, l_ in (("mpl", f"  其中可插拔（情境 {s_}）"), ("mcp", f"  其中 CPO（情境 {s_}）")):
            if r[0].value == l_:
                NETS[(k_, s_)] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETS) == 6 and all(v is not None for k in NETS for v in NETS[k][0]), NETS.keys()
NETR["其中可插拔"] = NETS[("mpl", "A")]; NETR["其中 CPO"] = NETS[("mcp", "A")]   # 情境 A 列（數值＝主表「其中可插拔」「其中 CPO」）
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。公司金額：十億美元（US$B；Q 頁百萬美元）；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司不揭露產品別營收（只有元件／系統兩個部門），產品拆分為 Assumed；口徑為公司 non-GAAP"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
PR = "https://www.sec.gov/Archives/edgar/data/0001633978/000162828026055726/lite_ex991xq4fy26.htm"
PPT = "https://s21.q4cdn.com/377324469/files/doc_financials/2026/q4/Q4-FY26-Earnings-Presentation_final.pdf"
K10 = "https://www.sec.gov/Archives/edgar/data/0001633978/000162828026057358/lite-20260627.htm"
MB = "https://www.marketbeat.com/instant-alerts/lumentum-q4-earnings-call-highlights-2026-08-11/"
ZK = "https://www.zacks.com/stock/news/2973530/lumentum-q4-earnings-call-focuses-on-ai-optics-ocs-npo-ramps"
KGI = "https://www.kgi.com.hk/en/-/media/files/kgishk/research-reports/us-daily/optical-networking-industry_20260818_tc.pdf"
TF = "https://www.trendforce.com/presscenter/news/20260603-13077.html"
CRUX = "https://cruxcapitalgroup.substack.com/p/lite-earnings-breakdown"
PC = "https://photoncap.net/p/the-8-companies-behind-lumentums"
SA = "https://stockanalysis.com/stocks/lite/forecast/"
TFL = "https://www.tickflow.io/stock/LITE/forecast"
TRE = "https://www.trefis.com/stock/lite/articles/617928/what-is-the-true-hidden-price-of-lumentum-stock/2026-10-08"
VCP = "https://www.vcpscanner.com/valuation/lite/estimates"
IVS = "https://www.investsmart.com.au/security/nasdaq/lite/lumentum-holdings/share-price"
AJ = "https://pdf.dfcfw.com/pdf/H3_AP202610121830372080_1.pdf"

sec("A. AI 光學市場與 TPU 顆數（讀 markets/network 成品；日曆年；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mpl_{y}", f"AI 可插拔光模組市場 CY20{y}", "US$B", None, NETR["其中可插拔"][0][j], None, "Derived", NETR["其中可插拔"][1],
        "加速器顆數 × 每顆光學內容 − CPO（網通模型；日曆年）" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mcp_{y}", f"AI CPO 光學市場 CY20{y}", "US$B", None, NETR["其中 CPO"][0][j], None, "Derived", NETR["其中 CPO"][1], "共同封裝光學；外部光源（ELS）與超高功率雷射屬此" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"tpu_{y}", f"Google TPU 顆數 CY20{y}", "百萬顆", None, NETR["TPU"][0][j], None, "Derived", NETR["TPU"][1], "OCS（光路交換）的主要用戶是 Google TPU 叢集" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 CY20{y}", "US$B", None, NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][0][j], None, "Derived", NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][1], "Recon 用：Lumentum AI 營收 ÷ AI IT 資本支出" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：CY2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、中際旭創模型相同）")
    inp(f"um29_{s}", f"情境 {s}：CY2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")

sec("B. 雷射晶片（EML／CW）：高速光模組數量 × 每模組雷射數 × 雷射單價 × Lumentum 份額（上游元件口徑，不可與光模組廠份額相加）")
for y, lo, b, hi in ((26, 420, 500, 600), (27, 440, 520, 640), (28, 450, 540, 680), (29, 450, 560, 720)):
    inp(f"asp_{y}", f"光模組平均單價 CY20{y}（把網通模型金額換成模組數）", "$/個", lo, b, hi, "Assumed",
        "800G 為主、1.6T 比重上升（每個模組單價上升）vs 每 Gb 降價（下壓）" if y == 26 else "同上",
        KGI + "：800G＋1.6T 出貨 2026 年 67.2M、2027 年 107.5M（凱基；只在 Recon 對照）" if y == 26 else "—")
for y, lo, b, hi in ((26, None, 5.5, None), (27, 4.5, 5.5, 6.0), (28, 4.0, 5.0, 6.0), (29, 3.5, 4.8, 6.0)):
    inp(f"lpm_{y}", f"每模組雷射數 CY20{y}（EML 與 CW 加權）", "顆", lo, b, hi, "Assumed",
        "EML 方案（DR8／2×FR4）每模組 8 顆；矽光方案每模組 1–4 顆高功率 CW（業界常識，未找到第一手來源）" if y == 26 else "矽光比重上升使每模組雷射數下降",
        PC + "：公司 FY2026 第四季約 20% 光模組改用自製 CW 雷射" if y == 26 else "—")
for y, lo, b, hi in ((26, None, 13, None), (27, 12, 15, 17), (28, 11, 15, 18), (29, 10, 14, 18)):
    inp(f"lasp_{y}", f"雷射平均單價 CY20{y}（EML 與 CW 加權）", "$/顆", lo, b, hi, "Assumed",
        AJ + "：光學元件與雷射約占光模組成本 25–35%（次級來源）" if y == 26 else "200G EML 比重上升（公司：2027 年中占 EML 出貨 50% 以上）",
        "雷射單價：未找到；本列×每模組雷射數＝每模組雷射內容（2026 約占模組單價 14%）" if y == 26 else "—")
inp("lrel_27", "雷射份額 CY2027 ÷ CY2026（相對值）", "x", 0.85, 0.95, 1.00, "Assumed", TF + "：EML 前三大（Lumentum、Broadcom、三菱電機）合計約 72%；CW 由 Broadcom、住友電工領先（TrendForce，Interested-party）", KGI + "：Lumentum 是唯一量產 200G EML 的供應商；EML 供需缺口約 30%（凱基）")
inp("lrel_28", "雷射份額 CY2028 ÷ CY2026", "x", 0.80, 0.95, 1.05, "Assumed", "Greensboro（第五座 InP 廠）2028 年初量產（公司）", "Coherent 擴 6 吋 InP、中國 EML 廠追趕（下壓）")
inp("lrel_29", "雷射份額 CY2029 ÷ CY2026", "x", 0.70, 0.90, 1.05, "Assumed", "同上", "—")

sec("C. 光模組、OCS、CPO（另推）")
inp("mrel_27", "光模組份額 CY2027 ÷ CY2026（相對值）", "x", 0.80, 1.0, 1.20, "Assumed", MB + "：800G 出貨創新高、1.6T 開始量產（公司）", "公司估光模組供給缺口約 30%（vocus 轉述法說，Interested-party）")
inp("mrel_28", "光模組份額 CY2028 ÷ CY2026", "x", 0.70, 1.0, 1.30, "Assumed", "中際旭創、新易盛、Coherent 為主要競爭者", "—")
inp("mrel_29", "光模組份額 CY2029 ÷ CY2026", "x", 0.60, 1.0, 1.30, "Assumed", "同上", "—")
inp("ocsk_27", "OCS 每顆 TPU 營收 CY2027 ÷ CY2026（相對值）", "x", 1.2, 1.6, 2.0, "Assumed", ZK + "：2H CY2026 OCS 約 4 億美元（追上原展望、未超前）；FY2027 第一季首度超過 1 億美元（公司）", PPT + "：多年期、數十億美元的 OCS 採購協議（公司）")
inp("ocsk_28", "OCS 每顆 TPU 營收 CY2028 ÷ CY2026", "x", 1.2, 1.8, 2.4, "Assumed", "2H CY2026 年化 8 億美元 ÷ CY2026 約 5 億美元 ≈ 1.6（推導）", KGI + "：Coherent 把 OCS 市場規模由 20 億上調到 40 億美元（Interested-party）")
inp("ocsk_29", "OCS 每顆 TPU 營收 CY2029 ÷ CY2026", "x", 1.0, 1.8, 2.4, "Assumed", "同上；Coherent 等競爭者加入（下壓）", "—")
inp("cpo_27", *SIL.row9("cposh_lite_27")[1:])
inp("cpo_28", *SIL.row9("cposh_lite_28")[1:])
inp("cpo_29", *SIL.row9("cposh_lite_29")[1:])

sec("D. CY2026 錨點的產品拆分（公司只揭露元件／系統；以下為 Assumed，雷射晶片＝殘差）")
inp("ocs_1h", "OCS 營收 1H CY2026（FY2026 第三、四季）", "US$B", 0.06, 0.10, 0.14, "Assumed", ZK + "：FY2026 第四季 OCS 出貨較第三季倍增（公司）", "金額：未找到")
inp("ocs_2h", "OCS 營收 2H CY2026（公司展望）", "US$B", None, 0.40, None, "Interested-party", ZK + "：約 4 億美元（公司）", KGI + "：同值（凱基轉述）")
inp("cpo_a26", "CPO／超高功率雷射營收 CY2026", "US$B", 0.03, 0.05, 0.08, "Interested-party", KGI + "：2026 年約 5 千萬美元（凱基轉述公司）", "—")
inp("tel_a26", "電信與 DCI 元件營收 CY2026（WSS、ROADM、泵浦、窄線寬雷射等）", "US$B", 0.55, 0.75, 0.95, "Assumed", MB + "：窄線寬雷射 +130%、泵浦雷射 +80%（年增；公司）", KGI + "：泵浦雷射份額 70–80%（凱基）")
inp("ind_a26", "工業雷射與 3D 感測營收 CY2026", "US$B", 0.25, 0.30, 0.40, "Assumed", "https://www.sec.gov/Archives/edgar/data/1633978/000162828024037191/lite_ex991xq4fy24.htm：FY2024 舊部門 Industrial Tech 2.743 億美元（公司）", "FY2026 金額：未找到")
inp("xcvr_a26", "光模組（收發器）營收 CY2026", "US$B", 0.70, 1.00, 1.30, "Assumed", PPT + "：系統部門 FY2026 第四季 3.569 億美元（含光模組、OCS、工業雷射）", "光模組金額：未找到；系統部門 CY2026 約 16 億 − OCS 約 5 億 − 工業雷射（推導）")
inp("g_tel_27", "電信與 DCI 營收年增 CY2027", "%", 0.10, 0.25, 0.40, "Assumed", "泵浦雷射出貨未來數季約 4 倍（公司）；DCI／scale-across", "—")
inp("g_tel_28", "電信與 DCI 營收年增 CY2028", "%", 0.0, 0.15, 0.30, "Assumed", "同上", "—")
inp("g_tel_29", "電信與 DCI 營收年增 CY2029", "%", -0.05, 0.10, 0.20, "Assumed", "同上", "—")
inp("g_ind", "工業雷射與 3D 感測年增（CY2027 起）", "%", -0.05, 0.03, 0.10, "Assumed", "FY2023 4.445 億 → FY2024 2.743 億美元（Industrial Tech；公司）", "非 AI；成長假設低")

sec("E. 季度實績與指引（百萬美元；CY2026 錨點）")
inp("q3_26", "FY2026 第三季營收（2026-01～03）", "US$M", None, 808.4, None, "Verified", PR + "：808.4", CRUX + "：同值")
inp("q4_26", "FY2026 第四季營收（2026-04～06）", "US$M", None, 1006.3, None, "Verified", PR + "：1,006.3（元件 649.4、系統 356.9）", PPT + "：同值")
inp("q1g", "FY2027 第一季營收指引中點（2026-07～09）", "US$M", None, 1250, None, "Interested-party", PR + "：12.25–12.75 億美元（公司）", KGI + "：同值")
inp("q_g", "FY2027 第二季營收季增（2026-10～12）", "%", 0.05, 0.12, 0.20, "Assumed", "FY2026 第三、四季、FY2027 第一季季增 21%、24%、24%（推導）", "產能受限：EML、泵浦缺口 30% 以上（公司）")
inp("om_g", "FY2027 第一季 non-GAAP 營業利益率指引中點", "%", None, 0.40, None, "Interested-party", PR + "：39.5–40.5%（公司）", "—")
inp("eps_g", "FY2027 第一季 non-GAAP EPS 指引中點", "US$", None, 4.20, None, "Interested-party", PR + "：4.05–4.35（公司）", "Zacks 共識 4.23（Recon 不用）")
inp("sh_q1", "FY2027 第一季 non-GAAP 稀釋股數指引", "百萬股", None, 102.0, None, "Interested-party", PPT + "：約 1.02 億股（含可轉債假設轉換）", K10 + "：普通股 8,970 萬股（2026-08-14）")
inp("tx", "non-GAAP 有效稅率", "%", 0.14, 0.165, 0.19, "Interested-party", PPT + "：16.5%（公司指引）", "—")
inp("h_1h", "日曆年上半年營收占全年比（CY2027 起；換算會計年度用）", "%", 0.44, 0.47, 0.50, "Assumed", "成長期上半年小於一半；CY2026 上半年占 40.6%（推導，高成長）", "會計年度 FY(t)＝CY(t−1) 下半年＋CY(t) 上半年")

sec("F. FY2025、FY2026 實績（non-GAAP；百萬美元）")
inp("a25_rev_in", "FY2025 營收", "US$M", None, 1645.0, None, "Verified", PR + "：1,645.0", "https://www.01net.it/lumentum-announces-fourth-quarter-and-full-fiscal-year-2026-results/：同值（轉載）")
inp("a25_gm_in", "FY2025 non-GAAP 毛利率", "%", None, 0.347, None, "Verified", PR + "：34.7%", "GAAP 28.0%")
inp("a25_opx_in", "FY2025 non-GAAP 營業費用（研發＋銷管）", "US$M", None, "=256.0+155.4", None, "Verified", PR + "：研發 256.0、銷管 155.4", "—")
inp("a25_ni_in", "FY2025 non-GAAP 淨利", "US$M", None, 146.4, None, "Verified", PR + "：146.4", "GAAP 25.9")
inp("a25_eps_in", "FY2025 non-GAAP 稀釋 EPS", "US$", None, 2.06, None, "Verified", PR + "：2.06", "GAAP 0.37")
inp("a26_rev_in", "FY2026 營收", "US$M", None, 3014.0, None, "Verified", PR + "：3,014.0（元件 2,005.6、系統 1,008.4）", "https://www.01net.it/lumentum-announces-fourth-quarter-and-full-fiscal-year-2026-results/：同值（轉載）")
inp("a26_gm_in", "FY2026 non-GAAP 毛利率", "%", None, 0.460, None, "Verified", PR + "：46.0%（第四季 50.4%）", "GAAP 41.7%")
inp("a26_opx_in", "FY2026 non-GAAP 營業費用（研發＋銷管）", "US$M", None, "=304.8+184.9", None, "Verified", PR + "：研發 304.8、銷管 184.9", "—")
inp("a26_ni_in", "FY2026 non-GAAP 淨利", "US$M", None, 782.3, None, "Verified", PR + "：782.3", "GAAP −6,935.1（含可轉債清償損失 7,756.6，非現金）")
inp("a26_eps_in", "FY2026 non-GAAP 稀釋 EPS", "US$", None, 8.67, None, "Verified", PR + "：8.67", "GAAP −92.96")

sec("G. 損益參數（FY2027 起）")
for y, lo, b, hi in ((27, 0.49, 0.52, 0.54), (28, 0.46, 0.51, 0.54), (29, 0.42, 0.49, 0.53)):
    inp(f"gm_{y}", f"non-GAAP 毛利率 FY20{y}", "%", lo, b, hi, "Assumed",
        PPT + "：FY2026 第四季 50.4%；FY2027 第一季指引隱含約 51–52%（營業利益率 40% ＋ 營業費用率約 11–12%，推導）" if y == 27 else ("200G EML、OCS 規模效益（上推）vs 光模組占比上升、降價（下壓）" if y == 28 else "供給缺口收斂後價格回落"),
        KGI + "：公司長期毛利率目標 42% 以上（凱基轉述）" if y == 27 else "—")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.10, "=Q!B26", 0.60, "Derived", "FY2026 營業費用年增 ÷ 營收年增（Q 頁）", "公司：營業利益率目標由 38–42% 可再上調 100–200bp（Zacks 轉述法說）")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.01, 0.0, 0.02, "Assumed", "無", "Sens 用 +2 個百分點")
inp("oi_oth", "營業外淨收入（利息等，FY2027 起每年）", "US$B", 0.0, "=4*(eps_g*sh_q1/(1-tx)-om_g*q1g)/1000", 0.08, "Derived", "由 FY2027 第一季指引推導：EPS × 股數 ÷（1 − 稅率）− 營業利益", "現金與短期投資 27.4 億美元（FY2026 年底，公司）")
inp("sh_g", "non-GAAP 稀釋股數年變動（FY2028 起）", "%", 0.0, 0.01, 0.03, "Assumed", "股票報酬稀釋；2032 年可轉債已計入稀釋股數（公司）", "A 系列可轉換特別股 290 萬股（2026-03 發行；公司）")
inp("mf", "AI 市場係數（Sens 用；CY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 CY2027 可插拔光模組 +84%", "±20%")

sec("H. 共識、券商、研調（只作 Recon 對帳，不作校準）")
inp("cons_rev27", "共識 FY2027 營收", "US$B", None, 6.3, None, "Interested-party", SA + "（23 家，2026-10-21）：63 億美元（56–72）", TRE + "：同值")
inp("cons_eps27", "共識 FY2027 EPS", "US$", None, 21.79, None, "Interested-party", SA + "（23 家，2026-10-21）：21.79（19.05–25.27）", TFL + "：16.10（12 家；該頁年度與季度數字不一致）")
inp("cons_eps28", "共識 FY2028 EPS（單一來源、5 家）", "US$", None, 23.60, None, "Interested-party", TFL + "（5 家，2026-10-24）：23.60（13.23–35.39）", "第二來源：未找到（stockanalysis 需付費）")
inp("cons_ni28", "共識 FY2028 淨利（trefis 轉述）", "US$B", None, 2.4, None, "Interested-party", TRE + "（2026-10-08）：約 24 億美元；FY2028 營收年增 55.5%", "口徑未說明")
for y, a, b in ((26, 54.2, 13.0), (27, 76.7, 30.8), (28, 90.9, 60.2)):
    inp(f"kgi_m{y}", f"凱基：800G＋1.6T 光模組出貨 20{y}", "百萬個", None, f"={a}+{b}", None, "Interested-party", KGI + f"：800G {a}M、1.6T {b}M", "券商；只作對照")
inp("tf_eml3", "EML 前三大合計份額（Lumentum、Broadcom、三菱電機）", "%", None, 0.72, None, "Interested-party", TF + "（TrendForce 2026-06-03）", "Lumentum 個別份額：未找到")
inp("cust_a", "最大客戶占 FY2026 營收", "%", None, 0.266, None, "Verified", K10 + "：客戶 A 26.6%、客戶 B 15.0%", "客戶名稱：未揭露")
inp("kgi_pe", "凱基：LITE 歷史平均 P/E", "x", None, 45.0, None, "Interested-party", KGI + "（2026-08-17）：歷史平均 45.0x、高點 66.2x、最新 42.7x", "計算口徑未說明；只作對照")

sec("I. 價格與估值")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 756.17, None, "Interested-party", VCP + "：756.17（2026-10-23）", IVS + "：756.24（頁面標 2026-10-24，應為 10-23 收盤）")
inp("px_hi", "2026 年內高點", "US$", None, 1143, None, "Interested-party", IVS + "：2026 年內低點 317.61、高點 1,143", "52 週高點：未找到（以年內高點代用）")
inp("sh_basic", "普通股股數（2026-08-14）", "百萬股", None, 89.7, None, "Verified", K10 + "：8,970 萬股", "—")
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")

sec("K. 目標價（DCF＋Comparable）與收斂法參數（v0.2 新增；黃底，Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；各市場同一值）", "DCF 與 Sens")
inp("tp_g", "長期成長率（DCF 永續；收斂法情境 B 的 FY2032 營收年增）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "—")
inp("tp_roe", "ROE（與 EPS 同口徑：non-GAAP 淨利 ÷ 平均股東權益，近三年平均；DCF 再投資率＝tp_g ÷ tp_roe）", "%", None, 0.2, None, "Assumed", "本模型沒有股東權益（資產負債表未建），non-GAAP ROE 算不出；依主串規則（2026-10-25）取 0.20", '參考（不採）：GAAP ROE stockanalysis FY2024 −47.25%、FY2025 2.48%、FY2026 −240.03%（FY2026 GAAP 淨損含可轉債公允價值損失）')
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（工作單 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "收斂法：情境 A 於 FY2032 的營收年增（＝tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法；FY2030 起由 FY2029 年增線性收斂", "FY2030 起為收斂法，非驅動因子推導")
inp("fade_gB", "收斂法：情境 B 於 FY2032 的營收年增（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "收斂法：情境 C 於 FY2032 的營收年增（＝tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "收斂法：情境 A 的中段淨利率（＝情境 A FY2029 淨利率）", "%", None, None, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "收斂法：情境 B 的中段淨利率（＝FY2029 與十年平均的中點；十年平均以 FY2029 × 0.8 代替）", "%", None, None, None, "Assumed", "十年平均淨利率（2016–2025，與本模型同口徑）：未找到，依工作單以 FY2029 × 0.8 代替；本列＝（FY2029＋FY2029 × 0.8）÷ 2", "Andy 可改")
inp("fade_mC", "收斂法：情境 C 的中段淨利率（＝十年平均；以情境 C FY2029 × 0.8 代替）", "%", None, None, None, "Assumed", "十年平均淨利率：未找到，依工作單以情境 C FY2029 淨利率 × 0.8 代替", "Andy 可改")
inp("cons_tp", "市場共識目標價（賣方平均）", 'US$', None, 1165.03, None, "Interested-party", 'stockanalysis 統計頁：27 位分析師平均 1,165.03（2026-10-23）', '第二來源：未找到')
inp("cons_feps", '共識 12 個月後前瞻 EPS（≈ FY2028；FY2029 共識未找到）', 'US$', None, '=cons_eps28', None, "Derived", 'tickflow FY2028 共識 23.60（5 家，In 頁 cons_eps28；單一來源，可信度低）；本模型前瞻 EPS 含 4／12 的 FY2029', "共識只放 Recon 與本列，不作校準")

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit); INROW[key] = r
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    if "未找到" in f"{s1}{s2}":   # 含「未找到」的列會進 HTML 清單：網址改寫成「網址見 Sources」（HTML 須自含、無外部網址）
        s1, s2 = (re.sub(r"https?://[^\s：；（）]+", "（網址見 Sources）", str(x)) for x in (s1, s2))
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 CY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與錨點（百萬美元；EPS 為美元；股數為百萬股）。藍字＝輸入；黑字＝公式。口徑：公司 non-GAAP（GAAP 含 FY2026 可轉債清償損失 77.6 億美元等非現金項目，不用於預測）"
Q["A2"] = "CY2026 錨點＝FY2026 第三季（2026-01～03）＋第四季（04～06）實績＋FY2027 第一季（07～09）指引中點＋第二季（10～12，季增假設）。FY2027 第一季財報（約 2026-11-05）公布後以實績取代第一季、並以新指引取代第二季"
hdr(Q, 3, ["項目", "FY2025A", "FY2026A", "FY26 Q3（實績）", "FY26 Q4（實績）", "FY27 Q1（指引）", "FY27 Q2（假設）", "CY2026 錨點", "2H CY2026", "標記", "說明"])
def qrow(r, label, vals, fmt="#,##0.0", tag="Verified", note=""):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
    Q.cell(r, 10, tag); Q.cell(r, 11, note)
qrow(4, "營收", {"B": "=a25_rev_in", "C": "=a26_rev_in", "D": "=q3_26", "E": "=q4_26", "F": "=q1g", "G": "=F4*(1+q_g)", "H": "=D4+E4+F4+G4", "I": "=F4+G4"},
     tag="Verified／Interested-party（指引）／Assumed（Q2）", note="FY(t) 由 CY(t−1) 7 月起算；CY2026＝FY2026 後兩季＋FY2027 前兩季")
qrow(5, "毛利率（non-GAAP）", {"B": "=a25_gm_in", "C": "=a26_gm_in"}, fmt="0.0%")
qrow(6, "毛利", {"B": "=B4*B5", "C": "=C4*C5"}, tag="Derived")
qrow(7, "營業費用（研發＋銷管，non-GAAP）", {"B": "=a25_opx_in", "C": "=a26_opx_in"})
qrow(8, "營業利益（non-GAAP）", {"B": "=B6-B7", "C": "=C6-C7", "F": "=F4*om_g"}, tag="Derived", note="FY2026 公司揭露營業利益率 29.8%；FY2027 第一季＝指引營業利益率 × 營收")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCF"}, fmt="0.0%", tag="Derived")
qrow(10, "淨利（non-GAAP）", {"B": "=a25_ni_in", "C": "=a26_ni_in", "F": "=eps_g*sh_q1"}, tag="Verified／Derived（Q1 指引）")
qrow(11, "淨利 ÷ 營業利益（營業外與稅）", {c: f"={c}10/{c}8" for c in "BCF"}, fmt="0.000", tag="Derived", note="FY2027 起＝（營業利益＋營業外）×（1 − 稅率）")
qrow(12, "稀釋股數（non-GAAP）", {"B": "=B10/B13", "C": "=C10/C13", "F": "=sh_q1"}, tag="Derived", note="淨利 ÷ EPS；公司揭露 FY2025 71.2、FY2026 90.2")
qrow(13, "EPS（non-GAAP）", {"B": "=a25_eps_in", "C": "=a26_eps_in", "F": "=eps_g"}, fmt="0.00")
qrow(14, "產品拆分：OCS", {"H": "=(ocs_1h+ocs_2h)*1000"}, tag="Assumed／Interested-party", note="2H CY2026 公司展望 4 億美元＋1H 假設")
qrow(15, "產品拆分：CPO／超高功率雷射", {"H": "=cpo_a26*1000"}, tag="Interested-party")
qrow(16, "產品拆分：電信與 DCI 元件", {"H": "=tel_a26*1000"}, tag="Assumed")
qrow(17, "產品拆分：工業雷射與 3D 感測", {"H": "=ind_a26*1000"}, tag="Assumed")
qrow(18, "產品拆分：光模組（收發器）", {"H": "=xcvr_a26*1000"}, tag="Assumed")
qrow(19, "產品拆分：雷射晶片（EML／CW；殘差）", {"H": "=H4-SUM(H14:H18)"}, tag="Derived（殘差）", note="元件部門（約 64.5%）≈ 雷射晶片＋電信與 DCI＋3D 感測＋CPO 雷射（Recon 對照）")
qrow(20, "AI 相關（雷射晶片＋光模組＋OCS＋CPO）÷ 營收", {"H": "=(H19+H18+H14+H15)/H4"}, fmt="0.0%", tag="Derived")
Q.cell(22, 1, "參考與推導（校準未使用共識）").font = BOLD
Q.cell(23, 1, "FY2026 營收年增"); Q["B23"] = "=C4/B4-1"
Q.cell(24, 1, "FY2026 營業費用年增"); Q["B24"] = "=C7/B7-1"
Q.cell(25, 1, "2H CY2026 ÷ CY2026（FY2027 換算用）"); Q["B25"] = "=I4/H4"
Q.cell(26, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B26"] = "=B24/B23"
Q.cell(27, 1, "元件部門占 FY2026 第四季營收（公司揭露 64.5%）"); Q["B27"] = "=649.4/1006.3"
Q.cell(28, 1, "1H CY2026 ÷ CY2026"); Q["B28"] = "=(D4+E4)/H4"
for rr in range(23, 29): Q[f"B{rr}"].number_format = "0.0%" if rr != 26 else "0.00"
for nm, rr in (("rev", 4), ("gm", 5), ("opx", 7), ("ni", 10), ("sh", 12), ("eps", 13)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$C${rr}")
for nm, rr in (("ocs", 14), ("cpo", 15), ("tel", 16), ("ind", 17), ("xcvr", 18), ("las", 19), ("rev", 4)):
    defname(f"c26_{nm}", f"Q!$H${rr}")
defname("w2h_26", "Q!$B$25")
Q.column_dimensions["A"].width = 46
for c in "BCDEFGHI": Q.column_dimensions[c].width = 13
Q.column_dimensions["J"].width = 26; Q.column_dimensions["K"].width = 80
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
CYS = (2026, 2027, 2028, 2029); CC = {2026: "C", 2027: "D", 2028: "E", 2029: "F"}
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
mk = lambda b, y, s: f"{b}_{y % 100}" if s == "A" else f"{b}{s}_{y % 100}"   # 網通 CPO 三情境：情境 A 讀 A 節，B、C 讀同字母列
DEFP = dict(scn="A", lrel="lrel", lasp="lasp", lpm="lpm", asp="asp", mrel="mrel", ocsk="ocsk", cpo="cpo", gm="gm", mf="1", h="h_1h", opxa="0", gtel="g_tel")
AI_KEYS = ("las", "xcvr", "ocs", "cpo")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    def row(key, label, unit, fs, fmt="#,##0.000", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t, head):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
        hdr(ws, r, head); r += 1
    c_ = lambda y: CC[y]
    F = CYS[1:]
    sec_("A. 日曆年產品營收（C 欄＝CY2026 錨點；D–F＝CY2027–CY2029）", ["項目", "單位", "CY2026（錨點）", "CY2027", "CY2028", "CY2029", "", "說明"])
    row("mpl", "AI 可插拔光模組市場 × 情境顆數係數", "US$B", {c_(y): f"={mk('mpl', y, s)}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in CYS})
    row("asp", "光模組平均單價", "$/個", {c_(y): f"={P['asp'] if y > 2026 else 'asp'}_{y % 100}" for y in CYS}, fmt="#,##0")
    row("mod", "高速光模組數量 ＝ 市場 ÷ 單價", "百萬個", {c_(y): f"={c_(y)}{R['mpl']}*1000/{c_(y)}{R['asp']}" for y in CYS}, fmt="#,##0.0")
    row("lpm", "每模組雷射數（EML／CW）", "顆", {c_(y): f"={P['lpm'] if y > 2026 else 'lpm'}_{y % 100}" for y in CYS}, fmt="0.0")
    row("lasp", "雷射平均單價", "$/顆", {c_(y): f"={P['lasp'] if y > 2026 else 'lasp'}_{y % 100}" for y in CYS}, fmt="0.0")
    row("ltam", "光模組雷射市場 ＝ 數量 × 每模組雷射數 × 單價", "US$B", {c_(y): f"={c_(y)}{R['mod']}*{c_(y)}{R['lpm']}*{c_(y)}{R['lasp']}/1000" for y in CYS})
    row("lsh0", "Lumentum 雷射有效份額（CY2026 校準）＝ 雷射晶片營收 ÷ 雷射市場", "%", {"C": f"=c26_las/1000/C{R['ltam']}"}, fmt="0.0%", note="Derived：雷射晶片營收為 CY2026 錨點殘差")
    row("lrel", "雷射份額相對值（÷ CY2026）", "x", {"C": "=1", **{c_(y): f"={P['lrel']}_{y % 100}" for y in F}}, fmt="0.00")
    row("lsh", "雷射有效份額", "%", {c_(y): f"=$C${R['lsh0']}*{c_(y)}{R['lrel']}" for y in CYS}, fmt="0.0%")
    row("las", "雷射晶片營收（EML／CW）", "US$B", {c_(y): f"={c_(y)}{R['ltam']}*{c_(y)}{R['lsh']}" for y in CYS})
    row("msh0", "光模組份額（CY2026 校準）＝ 光模組營收 ÷ 可插拔市場", "%", {"C": f"=c26_xcvr/1000/C{R['mpl']}"}, fmt="0.0%")
    row("msh", "光模組份額", "%", {"C": f"=C{R['msh0']}", **{c_(y): f"=$C${R['msh0']}*{P['mrel']}_{y % 100}" for y in F}}, fmt="0.0%")
    row("xcvr", "光模組營收", "US$B", {c_(y): f"={c_(y)}{R['mpl']}*{c_(y)}{R['msh']}" for y in CYS})
    row("tpu", "Google TPU 顆數 × 情境顆數係數", "百萬顆", {c_(y): f"=tpu_{y % 100}*{umx(y, s)}" for y in CYS}, fmt="#,##0.00")
    row("ocsp", "OCS 每顆 TPU 營收（CY2026 校準 × 相對值）", "$/顆", {"C": f"=c26_ocs/C{R['tpu']}", **{c_(y): f"=$C{r}*{P['ocsk']}_{y % 100}" for y in F}}, fmt="#,##0")
    row("ocs", "OCS 營收", "US$B", {c_(y): f"={c_(y)}{R['tpu']}*{c_(y)}{R['ocsp']}/1000" for y in CYS})
    row("mcp", "AI CPO 光學市場 × 情境顆數係數", "US$B", {c_(y): f"={mk('mcp', y, s)}*{umx(y, s)}" for y in CYS})
    row("cps", "Lumentum 占 CPO 光學金額", "%", {"C": f"=c26_cpo/1000/C{r - 1}", **{c_(y): f"={P['cpo']}_{y % 100}" for y in F}}, fmt="0.0%")
    row("cpo", "CPO／外部光源營收", "US$B", {c_(y): f"={c_(y)}{R['mcp']}*{c_(y)}{R['cps']}" for y in CYS})
    row("ai", "AI 相關營收（雷射晶片＋光模組＋OCS＋CPO）", "US$B", {c_(y): "=" + "+".join(f"{c_(y)}{R[k]}" for k in AI_KEYS) for y in CYS})
    row("tel", "電信與 DCI 元件", "US$B", {"C": "=c26_tel/1000", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gtel']}_{y % 100})" for y in F}})
    row("ind", "工業雷射與 3D 感測", "US$B", {"C": "=c26_ind/1000", **{c_(y): f"={c_(y - 1)}{r}*(1+g_ind)" for y in F}})
    row("crev", "營收（日曆年）", "US$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['tel']}+{c_(y)}{R['ind']}" for y in CYS})
    row("cg", "營收年增（日曆年）", "%", {c_(y): f"={c_(y)}{R['crev']}/{c_(y - 1)}{R['crev']}-1" for y in F}, fmt="0.0%")
    # 會計年度換算
    sec_("B. 會計年度損益（FY(t)＝CY(t−1) 下半年＋CY(t) 上半年；C、D 欄＝FY2025、FY2026 實績）", ["項目", "單位", "FY2025A", "FY2026A（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
    hh = P["h"]
    def fy(src, y):
        if y == 2027: return f"w2h_26*C{R[src]}+{hh}*D{R[src]}"
        a, b = {2028: ("D", "E"), 2029: ("E", "F")}[y]
        return f"(1-{hh})*{a}{R[src]}+{hh}*{b}{R[src]}"
    row("fai", "AI 相關營收（會計年度）", "US$B", {FC[y]: "=" + fy("ai", y) for y in (2027, 2028, 2029)}, note="FY2027＝CY2026 下半年（Q 頁）＋ h × CY2027")
    row("foth", "電信、DCI、工業與 3D 感測（會計年度）", "US$B", {FC[y]: f"={fy('tel', y)}+{fy('ind', y)}" for y in (2027, 2028, 2029)})
    row("rev", "營收", "US$B", {"C": "=a25_rev/1000", "D": "=a26_rev/1000", **{FC[y]: f"={FC[y]}{R['fai']}+{FC[y]}{R['foth']}" for y in (2027, 2028, 2029)}})
    row("g", "營收年增", "%", {FC[y]: f"={FC[y]}{R['rev']}/{FC[y - 1]}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI 相關 ÷ 營收", "%", {FC[y]: f"={FC[y]}{R['fai']}/{FC[y]}{R['rev']}" for y in (2027, 2028, 2029)}, fmt="0.0%")
    row("gm", "毛利率（non-GAAP）", "%", {"C": "=a25_gm", "D": "=a26_gm", **{FC[y]: f"={P['gm']}_{y % 100}" for y in (2027, 2028, 2029)}}, fmt="0.0%")
    row("gp", "毛利", "US$B", {FC[y]: f"={FC[y]}{R['rev']}*{FC[y]}{R['gm']}" for y in FYS})
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx/1000", "D": "=a26_opx/1000", **{FC[y]: f"={FC[y - 1]}{r}*(1+MAX(0,phi*{FC[y]}{R['g']}))" for y in (2027, 2028, 2029)}})
    row("opx", "營業費用（研發＋銷管）", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{FC[y]: f"={FC[y]}{R['opb']}+{P['opxa']}*{FC[y]}{R['rev']}" for y in (2027, 2028, 2029)}})
    row("oi", "營業利益", "US$B", {FC[y]: f"={FC[y]}{R['gp']}-{FC[y]}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {FC[y]: f"={FC[y]}{R['oi']}/{FC[y]}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("ni", "淨利（non-GAAP）", "US$B", {"C": "=a25_ni/1000", "D": "=a26_ni/1000", **{FC[y]: f"=({FC[y]}{R['oi']}+oi_oth)*(1-tx)" for y in (2027, 2028, 2029)}})
    row("nim", "淨利率", "%", {FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("sh_", "稀釋股數（non-GAAP）", "百萬股", {"C": "=a25_sh", "D": "=a26_sh", "E": "=sh_q1", **{FC[y]: f"={FC[y - 1]}{r}*(1+sh_g)" for y in (2028, 2029)}}, fmt="#,##0.0")
    row("eps", "EPS（non-GAAP）", "US$", {FC[y]: f"={FC[y]}{R['ni']}*1000/{FC[y]}{R['sh_']}" for y in FYS}, fmt="0.00")
    row("epsg", "EPS 年增", "%", {FC[y]: f"={FC[y]}{R['eps']}/{FC[y - 1]}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Lumentum 損益（US$B；non-GAAP）。FY2025A、FY2026A 實績，FY2027 起為模型（會計年度 6 月底結束；FY2027＝2026-07～2027-06）。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2027 三情境只差在網通 CPO 情境（A 替代快、C 慢），顆數係數從 CY2028 開始"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）；網通 CPO 情境 A（替代快）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85；網通 CPO 情境 B", "C": "情境 C：CY2028–29 顆數 ×0.6；網通 CPO 情境 C（替代慢）"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026A（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.00"), ("g", "  營收年增", "%", "0.0%"), ("fai", "  AI 相關營收", "US$B", "#,##0.00"), ("foth", "  電信、DCI、工業與 3D 感測", "US$B", "#,##0.00"),
                                ("aish", "  AI 相關 ÷ 營收", "%", "0.0%"), ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"), ("nim", "淨利率", "%", "0.0%"), ("ni", "淨利（non-GAAP）", "US$B", "#,##0.00"),
                                ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            if key in ("fai", "foth", "aish") and y < 2027: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：CY2026 模型營收 − 錨點；雷射晶片營收 − 殘差（校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：CY2026 營收差（US$B）"); PL[f"D{sr}"] = f"=C{PB['A']['crev']}-c26_rev/1000"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：CY2026 雷射晶片營收差（US$B）"); PL[f"D{sr}"] = f"=C{PB['A']['las']}-c26_las/1000"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：FY2027 上半年（Q 頁 2H CY2026）÷ 模型 FY2027 營收"); PL[f"D{sr}"] = f"=Q!I4/1000/E{PB['A']['rev']}"; PL[f"D{sr}"].number_format = "0.0%"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 60
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：數量與份額 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 高速光模組數量、每模組雷射數與 Lumentum 份額（日曆年；情境 A）。雷射份額是上游元件口徑（占光模組內雷射金額），不可與中際旭創等光模組廠份額相加"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("mod", "高速光模組數量（情境 A）", "百萬個", lambda y: f"=PL!{CC[y]}{A_['mod']}", "#,##0.0", "網通模型可插拔金額 ÷ 平均單價（Assumed）"),
        ("lpm", "每模組雷射數", "顆", lambda y: f"=PL!{CC[y]}{A_['lpm']}", "0.0", "EML 方案 8 顆、矽光 1–4 顆"),
        ("ltam", "光模組雷射市場", "US$B", lambda y: f"=PL!{CC[y]}{A_['ltam']}", "#,##0.00", ""),
        ("lsh", "Lumentum 雷射有效份額", "%", lambda y: f"=PL!{CC[y]}{A_['lsh']}", "0.0%", "CY2026 以錨點殘差校準"),
        ("las", "Lumentum 雷射晶片營收", "US$B", lambda y: f"=PL!{CC[y]}{A_['las']}", "#,##0.00", ""),
        ("msh", "Lumentum 光模組份額（模組廠口徑）", "%", lambda y: f"=PL!{CC[y]}{A_['msh']}", "0.0%", "與中際旭創份額同口徑"),
        ("ocsp", "OCS 每顆 TPU 營收", "$/顆", lambda y: f"=PL!{CC[y]}{A_['ocsp']}", "#,##0", ""),
        ("ocs", "OCS 營收", "US$B", lambda y: f"=PL!{CC[y]}{A_['ocs']}", "#,##0.00", ""),
        ("cps", "Lumentum 占 CPO 光學金額", "%", lambda y: f"=PL!{CC[y]}{A_['cps']}", "0.0%", ""),
        ("itr", "Lumentum AI 相關營收 ÷ AI IT 資本支出", "%", lambda y: f"=PL!{CC[y]}{A_['ai']}/itc_{y % 100}", "0.00%", "markets/network 的 AI IT 資本支出")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in CYS:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：雷射有效份額由 CY2026 產品拆分的殘差推得（拆分為 Assumed），是最需要每季檢查的輸入；每模組雷射數隨矽光比重上升而下降，單價隨 200G EML 比重上升而上升。").font = BOLD
UN.column_dimensions["A"].width = 56; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 50

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元；non-GAAP EPS）。目標價＝DCF 與 Comparable 加權（第 0 節）；舊法規則值情境加權保留參考；另列「現價 ÷ FY2027、FY2028 EPS（模型）」、市場隱含 EPS 與十年 P/E 規則值。十年 P/E 用年底股價 ÷ 近 12 月 non-GAAP EPS（推導；與模型 EPS 同口徑）"
V["A2"] = "循環位置：Lumentum 不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照中際旭創）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(14, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 15, ["項目", "A", "B", "C", "說明"])
V.cell(16, 1, "現價（2026-10-23 收盤）"); V["B16"] = "=px"; V["B16"].number_format = "0.00"
V.cell(17, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(18, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(19, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(20, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((17, 2026), (18, 2027), (19, 2028), (20, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(21, 1, "現價 ÷ 共識 FY2027 EPS"); V["B21"] = "=px/cons_eps27"; V["B21"].number_format = "0.0"
V.cell(22, 1, "現價 ÷ 2026 年內高點 − 1"); V["B22"] = "=px/px_hi-1"; V["B22"].number_format = "0%"
V.cell(23, 1, "市值（普通股 × 現價；US$B）"); V["B23"] = "=px*sh_basic/1000"; V["B23"].number_format = "#,##0.0"
V["E18"] = "FY2027 三情境只差在網通 CPO 情境（顆數係數從 CY2028 開始）"; V["E21"] = "共識 21.79（stockanalysis，23 家）"
V.cell(25, 1, "2. 循環位置（本模型營收最高的會計年度）").font = BOLD
hdr(V, 26, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 27):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(31, 1, "3. 規則值隱含股價（美元；約 FY2027 年底＝2027-06，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 32, ["情境", "FY2027 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 33 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=D{27 + j}").number_format = "0.0"
    V.cell(i, 4, f"=B{i}*C{i}").number_format = "#,##0"
    V.cell(i, 5, f"=D{i}/px-1").number_format = "0%"
    V.cell(i, 6, f"=px/B{i}").number_format = "0.0"
    for c in range(1, 7): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(37, 1, "4. 核心問題：市場隱含的 FY2028 EPS（＝現價 ÷ 假設倍數）與模型的距離").font = BOLD
hdr(V, 38, ["假設倍數（÷ FY2028 EPS）", "市場隱含 FY2028 EPS", "模型 FY2028 EPS（A）", "隱含 ÷ 模型 A − 1", "模型 FY2028 EPS（C）", "隱含 ÷ 模型 C − 1"])
vr = 39; IMPL = {}
for m in (15, 20, 30, "=pe_mid"):
    V.cell(vr, 1, m).number_format = "0.0"
    if not isinstance(m, str): V.cell(vr, 1).font = BLUE
    V.cell(vr, 2, f"=px/A{vr}").number_format = "0.00"
    V.cell(vr, 3, f"=PL!F{SUMR[('A', 'eps')]}").number_format = "0.00"
    V.cell(vr, 4, f"=B{vr}/C{vr}-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[('C', 'eps')]}").number_format = "0.00"
    V.cell(vr, 6, f"=B{vr}/E{vr}-1").number_format = "0%"
    IMPL[m] = vr; vr += 1
V.cell(vr, 1, "讀法：第四列＝十年 P/E 中位數（規則中段倍數）。若市場對 FY2028 只給歷史中位數倍數，現價隱含的 FY2028 EPS 遠高於模型——現價定價的是 AI 時代的較高倍數或比模型更高的獲利（見第 0 節目標價方法）。"); NOTE4 = vr
vr += 2
V.cell(vr, 1, "本版未做先驗：無 Andy 合理 P/E 對照。"); NOPRIOR = vr
assert vr < 50, vr
# 十年 P/E：來源 1（採用）＝年底股價（digrin）÷ 近 12 月 non-GAAP EPS（≈ 相鄰兩個會計年度 non-GAAP EPS 平均；公司新聞稿）
PXY = [38.65, 48.90, 42.01, 79.30, 94.80, 105.77, 52.17, 52.42, 83.95, 368.59]
EPSF = {2016: 1.29, 2017: 1.94, 2018: 3.82, 2019: 4.25, 2020: 4.95, 2021: 5.84, 2022: 6.05, 2023: 4.56, 2024: 1.01, 2025: 2.06, 2026: 8.67}
PE_NG = [f"={p}/(({EPSF[y]}+{EPSF[y + 1]})/2)" for p, y in zip(PXY, range(2016, 2026))]
PE_CMC = [161, 32.6, 27.5, None, 37.8, 20.5, 130, None, None, 103]   # companiesmarketcap GAAP 年底 P/E；2019、2023、2024 為負（虧損），不列
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "年底股價 ÷ 近 12 月 non-GAAP EPS（推導）", "vals": PE_NG, "tag": "Derived（股價：digrin 年底；EPS：公司新聞稿 FY2016–FY2026 non-GAAP 稀釋 EPS；近 12 月＝年底前後兩個會計年度平均）",
     "note": "FY2018 non-GAAP EPS 3.82 與 GAAP 同值（新聞稿擷取結果，待人工複核）；FY2020 起為公司新口徑"},
    {"label": "companiesmarketcap（GAAP 年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；GAAP 口徑含併購攤銷、可轉債損失，2019、2023、2024 為負不列；只作參考、不採用）", "use": False, "ref_rows": True}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=B{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=B{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=B{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝年底 P/E 依循環位置取分位。採用 non-GAAP 推導序列（與模型 EPS 同口徑）；GAAP 序列（companiesmarketcap）只列參考。2016–2023 年 Lumentum 以電信、3D 感測為主（AI 前），倍數約 10–24x；2024–2025 年 AI 轉型期倍數 55–69x。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
# 舊法：規則值情境加權（v0.1 稱「0. 目標價與上檔空間」；v0.2 改名保留，只作參考；建置後與其他內容一起下移）
V.cell(4, 1, "舊法：規則值情境加權（參考）").font = BOLD
hdr(V, 5, ["項目", "數值", "", "", "說明"])
V.cell(6, 1, "規則值情境加權價（參考，非目標價）"); V["B6"] = f"=tp_wA*D{VAL['A']}+tp_wB*D{VAL['B']}+tp_wC*D{VAL['C']}"; V["B6"].number_format = "#,##0"
V["E6"] = "權重 × 各情境隱含價（第 3 節規則值）"
V.cell(7, 1, "規則值情境加權價 ÷ 現價 − 1（參考）"); V["B7"] = "=B6/px-1"; V["B7"].number_format = "0.0%"; V["E7"] = "現價＝In 頁 px（2026-10-23 收盤）"
for rr, s, lab in ((8, "A", "情境 A 隱含價（樂觀）"), (9, "B", "情境 B 隱含價（基準）"), (10, "C", "情境 C 隱含價（下檔）")):
    V.cell(rr, 1, lab); V[f"B{rr}"] = f"=D{VAL[s]}"; V[f"B{rr}"].number_format = "#,##0"
    V[f"E{rr}"] = f'="FY2027 EPS "&TEXT(B{VAL[s]},"0.00")&" × "&TEXT(C{VAL[s]},"0.0")&"x（第 3 節；"&C{27 + SC.index(s)}&"）"'
V.cell(11, 1, "情境權重 A／B／C"); V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0%"
V["E11"] = '="合計 "&TEXT(B11+C11+D11,"0%")&"；In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"'
V.cell(12, 1, "舊法說明")
V["B12"] = ("P/E：各情境 FY2027（2026-07～2027-06）non-GAAP EPS × 十年（2016–2025 年底）P/E 依循環位置取分位（本模型營收到 FY2029 仍增加，三情境皆中段＝中位數）；"
            "P/E 為年底股價 ÷ 近 12 月 non-GAAP EPS（推導）。十年倍數 8 年落在 AI 前的電信、3D 感測期（約 10–24x），偏誤方向為低估；"
            "另 FY2027 EPS 三情境相近（只差在網通 CPO 情境）、且低於 FY2028，用 FY2027 也偏低。對照：凱基歷史平均 45x（口徑未說明）、現價 ÷ 共識 FY2027 EPS 約 35x")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；份額在每個案例內依 CY2026 錨點重算，CY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 相關營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 相關營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 相關營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("雷射份額流失（低）", "雷射份額相對值 CY2027 0.85、CY2028 0.80、CY2029 0.70（Broadcom、Coherent、中國 EML 追趕）", dict(lrel="LRLO")),
         ("雷射份額 高", "相對值 1.00／1.05／1.05（Greensboro 擴產）", dict(lrel="LRHI")),
         ("每模組雷射數 低（矽光替代）", "CY2027 4.5、CY2028 4.0、CY2029 3.5 顆", dict(lpm="LPLO")),
         ("雷射單價 低", "CY2027 $12、CY2028 $11、CY2029 $10", dict(lasp="LALO")), ("雷射單價 高", "$17／$18／$18", dict(lasp="LAHI")),
         ("光模組單價 高（模組數少）", "CY2027 $640、CY2028 $680、CY2029 $720", dict(asp="ASHI")),
         ("光模組份額 低", "相對值 0.80／0.70／0.60", dict(mrel="MRLO")), ("光模組份額 高", "相對值 1.20／1.30／1.30", dict(mrel="MRHI")),
         ("OCS 低", "每顆 TPU 營收相對值 1.2／1.2／1.0", dict(ocsk="OKLO")), ("OCS 高", "2.0／2.4／2.4", dict(ocsk="OKHI")),
         ("CPO 份額 低", "10%", dict(cpo="CPLO")), ("CPO 份額 高", "35%／35%／30%", dict(cpo="CPHI")),
         ("毛利率 低", "FY2027 49%、FY2028 46%、FY2029 42%", dict(gm="GMLO")), ("毛利率 高", "54%／54%／53%", dict(gm="GMHI")),
         ("AI 市場 ×0.8", "CY2027 起可插拔光模組市場 ×0.8", dict(mf="mf_lo")), ("AI 市場 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("上半年占比 44%", "日曆年上半年占 44%（會計年度換算）", dict(h="h_1h_lo")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi"))]
for nm, base, suf in (("LRLO", "lrel", "_lo"), ("LRHI", "lrel", "_hi"), ("LPLO", "lpm", "_lo"), ("LALO", "lasp", "_lo"), ("LAHI", "lasp", "_hi"), ("ASHI", "asp", "_hi"),
                      ("MRLO", "mrel", "_lo"), ("MRHI", "mrel", "_hi"), ("OKLO", "ocsk", "_lo"), ("OKHI", "ocsk", "_hi"), ("CPLO", "cpo", "_lo"), ("CPHI", "cpo", "_hi"),
                      ("GMLO", "gm", "_lo"), ("GMHI", "gm", "_hi")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}_{y}"] = DefinedName(f"{nm}_{y}", attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.00"), ("D", "fai", 2027, "#,##0.00"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.00"), ("I", "fai", 2028, "#,##0.00"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85；CPO 情境 B", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6；CPO 情境 C", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2027 有一半來自 CY2026 下半年（指引＋假設），所以 FY2027 對各輸入的敏感度約只有 FY2028 的一半；顆數情境 B、C 的差異從 FY2028 起。雷射份額、單價、每模組雷射數三者相乘，是雷射晶片營收的三條路徑。")
SENS_NOTE = nn + 3
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、共識、券商與研調。只作事後對照，不為對上數字而調參數。共識區間寬（FY2027 EPS 19.05–25.27），FY2028 共識只有單一來源、5 家"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
cy = lambda key, y: (lambda s: f"=PL!{CC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與券商（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "stockanalysis（23 家；56–72 億美元）", "FY2027 上半年＝指引＋季增假設")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "stockanalysis（23 家；19.05–25.27）", "")
rec("FY2028 EPS vs 共識（單一來源）", "=cons_eps28", pl("eps", 2028), "tickflow（5 家；13.23–35.39）", "該頁年度與季度數字不一致；可信度低")
rec("FY2028 淨利 vs trefis 轉述共識", "=cons_ni28", pl("ni", 2028), "trefis（2026-10-08）", "口徑未說明")
RC.cell(rr_, 1, "B. 數量、份額與產品（機制檢查）").font = BOLD; rr_ += 1
for y in (2026, 2027, 2028):
    rec(f"CY{y} 高速光模組數量 vs 凱基 800G＋1.6T 出貨（百萬個）", f"=kgi_m{y % 100}", cy("mod", y), "凱基 2026-08-18", "差距＝網通模型金額 ÷ 本模型單價 vs 券商數量；凱基 2028 年遠高於網通模型（網通模型 2028 年保守）", fmt="#,##0.0")
rec("CY2026 雷射有效份額 vs EML 前三大合計 72%（百分點）", "=tf_eml3", cy("lsh", 2026), "TrendForce", "有效份額含 CW、以金額計；EML 前三大合計的三分之一約 24%，Lumentum 為 EML 第一、200G 唯一量產", fmt="0.0%", pt=True)
rec("CY2026 雷射晶片 ÷ 營收", None, lambda s: f"=PL!C{PB[s]['las']}/(c26_rev/1000)", "Q 頁殘差", "", fmt="0.0%")
rec("CY2026 元件類（雷射＋電信與 DCI＋CPO＋3D 感測一半）÷ 營收 vs 元件部門 64.5%（百分點）", "=Q!B27", lambda s: f"=(PL!C{PB[s]['las']}+PL!C{PB[s]['tel']}+PL!C{PB[s]['cpo']}+PL!C{PB[s]['ind']}/2)/(c26_rev/1000)", "公司 FY2026 第四季部門", "3D 感測（元件）與工業雷射（系統）各半，Assumed", fmt="0.0%", pt=True)
rec("CY2027 OCS 營收 vs Coherent 估 OCS 市場 40 億美元（占比）", None, lambda s: f"=PL!D{PB[s]['ocs']}/4", "凱基轉述 Coherent", "", fmt="0.0%")
rec("CY2027 AI 相關營收 ÷ AI IT 資本支出", None, lambda s: f"=PL!D{PB[s]['ai']}/itc_27", "markets/network", "每 1 美元 AI IT 資本支出分到的 Lumentum 營收", fmt="0.00%")
rec("CY2028 雷射有效份額", None, cy("lsh", 2028), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("FY2025 EPS（淨利 ÷ 推導股數）vs 公司揭露（股數由 EPS 推導，恆等）", "=a25_eps_in", lambda s: "=Q!B13", "公司", "公司揭露稀釋股數 71.2（推導 71.1）", fmt="0.00")
rec("FY2026 EPS vs 公司揭露 8.67", "=a26_eps_in", pl("eps", 2026), "公司", "公司揭露稀釋股數 90.2", fmt="0.00")
rec("FY2027 第一季營業外淨收入（推導，年化 US$B）", None, lambda s: "=oi_oth", "Q1 指引", "", fmt="0.000")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 Lumentum AI 營收 ÷ 部署 GW 計算每 GW 營收"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.cell(rr_, 1, "FY2027 第一季財報（約 2026-11-05）：公布後以實績取代 Q 頁第一季、以新指引取代第二季"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_END = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝SEC 8-K 新聞稿、10-K、法說簡報；法說內容經第三方轉述者待人工複核）"; SO["A1"].font = BOLD
srcs = [("公司：FY2026 第四季與全年新聞稿（SEC 8-K EX-99.1）", PR),
        ("公司：FY2026 第四季法說簡報", PPT),
        ("公司：FY2026 10-K（部門定義、客戶集中度、股數）", K10),
        ("公司：歷年新聞稿（FY2016–FY2024 non-GAAP EPS；SEC 8-K）", "https://www.sec.gov/Archives/edgar/data/1633978/000162828024037191/lite_ex991xq4fy24.htm；…/000162828022022980/lite_ex991xq4fy22.htm；…/000162828021017219/lite_ex991xq4fy21.htm；…/000163397819000068/lite_ex991xq4fy19.htm；…/000163397818000056/lite_ex991xq4fy18.htm；…/000163397817000058/lite_ex991xq4fy17.htm"),
        ("二手：法說重點（MarketBeat、Zacks、Crux Capital、PhotonCap、vocus）", MB + "；" + ZK + "；" + CRUX + "；" + PC),
        ("券商：凱基 2026-08-18 光通訊產業報告（OCS、超高功率雷射、800G／1.6T 出貨、歷史 P/E）", KGI),
        ("研調：TrendForce 2026-06-03（EML、CW 產能與份額）", TF),
        ("券商：愛建證券 10-12 通信行業報告（光模組成本結構）", AJ),
        ("共識：stockanalysis（2026-10-21）、tickflow（2026-10-24）、trefis（2026-10-08）", SA + "；" + TFL + "；" + TRE),
        ("價格：vcpscanner（10-23）、investsmart（年內高低點）、digrin（年底股價）", VCP + "；" + IVS + "；https://www.digrin.com/stocks/detail/LITE/price"),
        ("資料商：companiesmarketcap GAAP 年底 P/E", "https://companiesmarketcap.com/lumentum/pe-ratio/"),
        ("市場層：markets/network、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：FY2027 第一季財報（約 11-05）；產品別營收（雷射、光模組、OCS）；Lumentum 個別 EML／CW 份額；雷射單價；FY2028 共識第二來源；52 週高點；FY2018 non-GAAP EPS 複核", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Lumentum（LITE）收支與估值模型 v0.2.2（AI 半導體 Project：網通次產業內複製——雷射與光元件）　2026-10-26", True),
 ("版本：v0.2.1（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）；v0.2（2026-10-25：三情境延伸到 FY2032、目標價＝DCF 與 Comparable 加權；v0.1 初版）。Lumentum 的雷射有效份額（上游元件口徑）在本模型 Units 頁建立；不可與中際旭創等光模組廠份額相加。", False),
 ("模型命題：Lumentum 的雷射晶片營收＝高速光模組數量（本專案網通模型金額 ÷ 模組單價）× 每模組 EML／CW 雷射數 × 雷射單價 × Lumentum 份額；光模組、OCS（Google TPU）、CPO 外部光源另推；電信與 DCI、工業與 3D 感測依成長率推。股價約 35 倍共識 FY2027 EPS，遠高於十年 non-GAAP P/E 中位數。", False),
 ("核心問題：EML 供給缺口收斂、矽光以較少 CW 雷射替代、OCS 與 CPO 放量時程——哪一條路徑決定 FY2028 獲利，現價隱含的倍數與獲利要多少才合理？", False), ("", False),
 ("會計年度 6 月底結束：FY2027＝2026-07～2027-06（FY(t) 後半年在日曆年 t）。市場層為日曆年，先算 CY2026–CY2029 產品營收，再換成 FY2027–FY2029（FY(t)＝CY(t−1) 下半年＋CY(t) 上半年，上半年占比 h＝47%）。口徑：公司 non-GAAP；EPS 以 non-GAAP 稀釋股數（含可轉債假設轉換）。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("雷射晶片：markets/network 可插拔光模組市場 × 情境顆數係數 ÷ 模組單價 × 每模組雷射數 × 雷射單價 × 有效份額（CY2026 以產品拆分殘差校準）× 份額相對值。光模組：可插拔市場 × 份額。OCS：Google TPU 顆數（網通模型 Units）× 每顆 TPU 營收。CPO：CPO 光學市場 × 份額。", False),
 ("CY2026 錨點＝FY2026 第三、四季實績＋FY2027 第一季指引中點＋第二季季增 12%（Assumed）。產品拆分（OCS、CPO、電信與 DCI、工業與 3D 感測、光模組）為 Assumed／公司展望，雷射晶片＝殘差。毛利率 FY2027 起為輸入路徑；營業費用彈性 phi（Derived）；稅率 16.5%；營業外淨收入由第一季指引推導。", False),
 ("估值：目標價（Val 第 0 節）＝0.5 × DCF（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，情境權重 0.3／0.5／0.2）＋0.5 × Comparable（美股光通訊同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）。舊法規則值情境加權保留參考（不是目標價）；規則值＝FY2027 EPS × 十年 P/E 依循環位置取分位（年底股價 ÷ 近 12 月 non-GAAP EPS，推導）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司不揭露產品別營收；法說內容多經第三方轉述。", False),
 ("", False),
 ("已知限制：(1) 產品拆分、雷射單價、每模組雷射數皆為 Assumed，雷射份額由殘差推得；(2) CY2026 下半年第二季為假設，FY2027 第一季財報約 11-05 公布；(3) 網通模型 CY2028 可插拔光模組只 +12%，比券商出貨預測保守；(4) 未設 EML 產能上限（公司 EML 出貨年增約 50%）；(5) 十年 P/E 為推導序列、多數年份在 AI 前，規則值偏低；(6) R4 產業部署路徑未建（Recon 預留列）。", False)]
lines.append(("v0.2（2026-10-25）：Val 頁最前面新增「0. 目標價」（0.1 DCF、0.2 Comparable、0.3 結論），既有內容整體下移；原「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留。PL 三情境延伸到 FY2032（FY2030 起為收斂法，非驅動因子推導；灰底），說明欄移到 K 欄。In 頁新增 K 節（tp_r、tp_g、tp_roe、tp_w_dcf、收斂法參數、共識目標價）。Sens 加目標價參數列。FY2026–FY2029 既有數字不變。", False))
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170


# ================= v0.2（2026-10-25）：FY2030–FY2032 收斂法；目標價＝DCF 與 Comparable 加權 =================
# 工作單：Andy 2026-10-25 11:18／11:20（三情境延伸到 FY2032、收斂法）、11:28（目標價＝DCF＋Comparable）、11:53（tp_w_dcf 0.5；A 股與其他市場同業分開）
# 本節在其他工作表都建好之後執行：(1) PL 說明欄由 H 移到 K，H–J 放 FY2030–32 收斂年度；(2) Val 既有內容整體下移，最前面放「0. 目標價」；(3) Sens 加目標價參數列
import re as _re
from copy import copy as _copy
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
FCX = dict(FC); FCX.update({2030: "H", 2031: "I", 2032: "J"})
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
INR = {IN.cell(i, 1).value: i for i in range(4, IN.max_row + 1) if IN.cell(i, 1).value}
for _k, _i in INR.items():
    if str(_k).startswith(("tp_", "fade_")) or _k in ("cons_tp", "cons_feps"):
        for _j in range(1, 10): IN.cell(_i, _j).fill = KEY

def _move_notes(ws):
    """說明欄（第 8 欄，H）移到第 11 欄（K），讓出 H–J；表頭列補上收斂年度。"""
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 8)
        if c.value is None: continue
        assert not (isinstance(c.value, str) and c.value.startswith("=")), (ws.title, r, c.value)
        k = ws.cell(r, 11); k.value = c.value; k.font = _copy(c.font); k.fill = _copy(c.fill)
        c.value = None; c.font = Font(); c.fill = PatternFill()
        if str(ws.cell(r, 1).value or "") in ("項目", "情境／項目"):
            for j, y in zip((8, 9, 10), FADE):
                h = ws.cell(r, j, f"FY{y}（收斂）"); h.font = BOLD; h.fill = HDR
    ws.column_dimensions["K"].width = ws.column_dimensions["H"].width or 70
    for col in "HIJ": ws.column_dimensions[col].width = 14

def _fade(ws, R, s, eps_f):
    """FY2030–32：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到）；非驅動因子推導。"""
    for y in FADE:
        c, pc, k = FCX[y], FCX[y - 1], y - 2029
        fs = {"g": f"=$G{R['g']}+(fade_g{s}-$G{R['g']})*{k}/3", "rev": f"={pc}{R['rev']}*(1+{c}{R['g']})",
              "nim": f"=$G{R['nim']}+(fade_m{s}-$G{R['nim']})*{k}/3", "ni": f"={c}{R['rev']}*{c}{R['nim']}",
              "sh_": f"={pc}{R['sh_']}*(1+sh_g)", "eps": "=" + eps_f(c, R), "epsg": f"={c}{R['eps']}/{pc}{R['eps']}-1"}
        for key, f in fs.items():
            ws[f"{c}{R[key]}"] = f; ws[f"{c}{R[key]}"].number_format = ws[f"G{R[key]}"].number_format; ws[f"{c}{R[key]}"].fill = FADEFILL

def _shift_rows(wb, sheet, start, n):
    """把 sheet 第 start 列起的內容整體下移 n 列，並同步改寫所有公式與定義名稱中指向這些列的參照（等同插入列）。"""
    ws = wb[sheet]
    cell_ = r"\$?[A-Z]{1,3}\$?\d+"
    pat = _re.compile(r"(?<![A-Za-z0-9_.!$:])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(" + cell_ + r")(?::(" + cell_ + r"))?(?![A-Za-z0-9_(!])")
    one = _re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")
    def fix(f, own):
        parts = f.split('"')
        for i in range(0, len(parts), 2):
            def rep(m):
                pre = m.group(1)
                tgt = (pre is None and own) or (pre is not None and pre.strip("'!") == sheet)
                def sh(x):
                    a, rr = one.fullmatch(x).groups()
                    rr = int(rr)
                    return f"{a}{rr + n if (tgt and rr >= start) else rr}"
                return (pre or "") + sh(m.group(2)) + (":" + sh(m.group(3)) if m.group(3) else "")
            parts[i] = pat.sub(rep, parts[i])
        return '"'.join(parts)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = fix(c.value, w.title == sheet)
    for nm in list(wb.defined_names.keys()):
        d = wb.defined_names[nm]
        if sheet + "!" in d.attr_text: d.attr_text = fix(d.attr_text, False)
    cells = [(c.row, c.column) for row in ws.iter_rows(min_row=start) for c in row if c.value is not None or c.has_style]
    for r, col in sorted(cells, reverse=True):
        src = ws.cell(r, col); dst = ws.cell(r + n, col)
        dst.value = src.value; dst._style = _copy(src._style); dst.number_format = src.number_format
        src.value = None; src._style = _copy(ws.cell(1, 30)._style)

def _tp_section(V, eps_ref, fwd, fwd_note, peers, peer_note, cons_note, method, pe_rng, fmt, cons_feps_note):
    """Val 第 0 節（第 4–48 列）：0.1 DCF、0.2 Comparable、0.3 結論。列標籤依工作單逐字；B／C／D＝情境 A 順利／B 平平／C 不如預期。"""
    T = {}
    def tset(r, key, label, fs, nf, note=""):
        V.cell(r, 1, label); T[key] = r
        for c, f in zip("BCD", fs):
            if f is None: continue
            V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = nf
        if note: V[f"E{r}"] = note
    V.cell(4, 1, "0. 目標價").font = BOLD
    V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
    hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
    V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法（PL 頁灰底）"
    for i, y in enumerate(YRS6):
        tset(8 + i, f"eps{y}", f"  FY{y}", [f"={eps_ref(s, y)}" for s in SC], "0.00")
    V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）")
    tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=IF(tp_roe<=tp_g,0.8,MIN(tp_g/tp_roe,0.8))", None, None], "0.0%", "ROE＝In 頁 tp_roe；tp_roe ≤ tp_g 時取上限 0.8")
    for i, y in enumerate(YRS6):
        tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{8 + i}*(1-$B$15)" for c in "BCD"], "0.00")
    tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], fmt,
         "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027")
    tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], fmt)
    tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], fmt, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
    V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
    hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
    assert 3 <= len(peers) <= 5, peers
    for i, (nm, tk, pe, dt, src) in enumerate(peers, 28):
        V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
    tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{27 + len(peers)})", None, None], "0.0", peer_note)
    tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）",
         ["=" + "+".join(f"tp_w{s}*(" + "+".join(f"{w}*{c}{8 + YRS6.index(y)}" for y, w in fwd) + ")" for s, c in zip(SC, "BCD")), None, None], "0.00", fwd_note)
    tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], fmt, "同業中位數 × 前瞻 EPS")
    V.cell(37, 1, "0.3 結論").font = BOLD
    hdr(V, 38, ["項目", "數值", "", "", "說明"])
    tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], fmt, "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
    tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], fmt)
    tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], fmt, cons_note)
    tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px")
    tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
    tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", cons_feps_note)
    tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
    V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
    tset(46, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D23,"{fmt}")&"～"&TEXT(B23,"{fmt}")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
    tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
         [f"=IF(px/{c}9>MAX({pe_rng}),1,IF(px/{c}9<MIN({pe_rng}),0,PERCENTRANK({pe_rng},px/{c}9)))" for c in "BCD"], "0%",
         "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 6 節採用來源）的分位；高於十年最大值＝100%；只作檢查，不進目標價")
    V.cell(48, 1, "目標價方法"); V["B48"] = method
    for r_ in (24, 35, 39):
        V[f"B{r_}"].fill = KEY
    return T

def _tp_sens(SE, T):
    """Sens：目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.3／0.7）；放在案例摘要與第 40 列推導區塊之間。"""
    last = max(r for r in range(1, 40) if any(SE.cell(r, j).value is not None for j in range(1, 13)))
    t0 = last + 2
    SE.cell(t0, 1, "目標價參數（只改 DCF 參數與權重；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
    hdr(SE, t0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
    def dcf_f(c, r, g, reinv):
        cf = lambda i: f"Val!{c}{8 + i}*(1-{reinv})"
        pv = "+".join(f"{cf(i)}/(1+{r})^{i + 1}" for i in range(6))
        return f"({pv}+{cf(5)}*(1+{g})/({r}-{g})/(1+{r})^6)"
    cases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
             ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
             ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
    for k, (nm, st, kw) in enumerate(cases):
        i = t0 + 2 + k; r = kw.get("r", "tp_r"); g = kw.get("g", "tp_g"); w = kw.get("w", "tp_w_dcf")
        reinv = f"IF(tp_roe<={g},0.8,MIN({g}/tp_roe,0.8))"
        SE.cell(i, 1, nm); SE.cell(i, 2, st)
        SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_f(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
        SE[f"D{i}"] = f"=Val!B{T['comptp']}"
        SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
        SE[f"F{i}"] = f"=E{i}/E${t0 + 2}-1"
        for c in "CDE": SE[f"{c}{i}"].number_format = "0.00"
        SE[f"F{i}"].number_format = "0.0%"
    assert t0 + 1 + len(cases) < 40, t0
    return t0, t0 + 1 + len(cases)

def _fade_all(eps_f, nim10):
    """PL：三情境推導區塊與頂端摘要延伸到 FY2032；In 頁 fade_m 連到 PL 頁 FY2029 淨利率。nim10＝十年平均淨利率的名稱；None＝未找到，依工作單用 FY2029 × 0.8。"""
    _move_notes(PL)
    for s in SC:
        _fade(PL, PB[s], s, eps_f)
        for key in ("rev", "g", "nim", "ni", "eps", "epsg"):
            if (s, key) not in SUMR: continue
            for y in FADE:
                cc = PL[f"{FCX[y]}{SUMR[(s, key)]}"]; cc.value = f"={FCX[y]}{PB[s][key]}"
                cc.number_format = PL[f"G{SUMR[(s, key)]}"].number_format; cc.fill = FADEFILL
    nimr = {s: f"PL!$G${PB[s]['nim']}" for s in SC}
    IN.cell(INR["fade_mA"], 5, f"={nimr['A']}")
    IN.cell(INR["fade_mB"], 5, f"=({nimr['B']}+{nim10})/2" if nim10 else f"=({nimr['B']}+{nimr['B']}*0.8)/2")
    IN.cell(INR["fade_mC"], 5, f"={nim10}" if nim10 else f"={nimr['C']}*0.8")
    PL["A2"] = ("灰底＝FY2030–FY2032 收斂法，非驅動因子推導（上游市場層只到 2029）：營收年增由 FY2029 線性收斂到 fade_g（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt），"
                "淨利率由 FY2029 線性收斂到 fade_m（In 頁 K 節），FY2032 達到；股數依 sh_g")

def _rule_pe_range(V, ncols=1):
    """第 6 節十年倍數表（2016–2025 列）中採用來源的欄位範圍（PERCENTRANK 用；ncols＞1 時各來源合併成一個分布）。"""
    yr = [r for r in range(1, V.max_row + 1) if V.cell(r, 1).value in [str(y) for y in range(2016, 2026)]]
    yr = [r for r in yr if r > 50][:10]
    assert len(yr) == 10 and yr[-1] - yr[0] == 9, yr
    return f"$B${yr[0]}:${chr(65 + ncols)}${yr[-1]}"

_fade_all(lambda c, R: f"{c}{R['ni']}*1000/{c}{R['sh_']}", None)
_shift_rows(wb, "Val", 4, 47)
TPT = _tp_section(V, lambda s_, y: f"PL!{FCX[y]}{SUMR[(s_, 'eps')]}", [(2028, '8/12'), (2029, '4/12')],
    '約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 8／12 × FY2028＋4／12 × FY2029（會計年度 6 月底結束）',
    [('Coherent', 'COHR', 22.3, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'), ('Fabrinet', 'FN', 17.27, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'), ('Applied Optoelectronics', 'AAOI', 29.34, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'), ('Ciena', 'CIEN', 30.77, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'), ('Corning', 'GLW', 32.89, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）')],
    '美股光通訊同業（不放 A 股）',
    'stockanalysis：27 位分析師平均 1,165.03（2026-10-23）',
    'DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%；ROE 需與 non-GAAP EPS 同口徑，本模型無股東權益、算不出，取 0.20（Assumed；再投資率 20%），FY2027–32 明確期＋永續，FY2030 起為收斂法。Comparable：美股光通訊同業（Coherent、Fabrinet、AAOI、Ciena、Corning）前瞻 P/E 中位數 × 12 個月後前瞻 EPS（non-GAAP）。會計年度 6 月底結束：DCF 的 t＝1 對 FY2027（2026-07～2027-06），偏高估。情境對應：顆數情境 A＝順利、B＝平平、C＝不如預期。',
    _rule_pe_range(V, 1), '#,##0.00',
    "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝In 頁 cons_feps")
TPS0, TPS_LAST = _tp_sens(SE, TPT)

# ---------------- v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 14:00、14:17；既有數值與儲存格位置不變） ----------------
_UPW = [r_ for r_ in range(1, 60) if V.cell(r_, 1).value == "上檔空間：加權目標價"]
assert _UPW == [42], _UPW
GR_ROW, RC_ROW = 49, 50   # 0.3 結論後的空白預留列（不移動其他列）
for _r in (GR_ROW, RC_ROW):
    assert all(V.cell(_r, _c).value is None for _c in range(1, V.max_column + 1)), _r
V.cell(GR_ROW, 1, "上檔空間級距")
V.cell(GR_ROW, 2, '=IF(B42>0.5,"大於 50%",IF(B42>0.25,"25%–50%",IF(B42>=0,"0%–25%",IF(B42>=-0.25,"0%～−25%",IF(B42>=-0.5,"−25%～−50%","小於 −50%")))))').number_format = "@"
V.cell(GR_ROW, 5, "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
V.cell(RC_ROW, 1, "投資建議")
V.cell(RC_ROW, 2, '=IF(B42>0.5,"強力買進",IF(B42>0.25,"買進",IF(B42>=-0.25,"中立",IF(B42>=-0.5,"放空","強力放空"))))').number_format = "@"
V.cell(RC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = [
       "目標價 US$623（DCF 387、Comparable 860 各半），比現價 756.17（2026-10-23）低 17.6%，級距 0%～−25%；Comparable 單獨為 +13.7%。DCF 拉低：DCF 單獨 −48.8%（ROE 用預設 20%），而本模型 12 個月後前瞻 EPS 29.30 比共識前瞻 EPS 23.60 高 24.2%（共識 FY2028 只有單一來源）；共識目標價 1,165.03 隱含 49.4 倍，比同業中位數 29.3 倍高 68.3%，市場給 EML／CW 雷射供給缺口的溢價，Comparable 未計入。",
       "最敏感的是 AI 市場與毛利率：AI 市場 ×0.8 時 FY2028 EPS −12.7%、毛利率低（FY2028 46%）−11.3%、雷射單價低 −9.5%、每模組雷射數低（矽光替代）−7.9%；情境 C FY2028 EPS 25.04，比情境 A 的 32.73 低 23.5%。現價 ÷ FY2028 EPS 23.1～30.2 倍，在十年分位 76%～80%。",
       "追蹤 FY2027 Q1 財報（約 2026-11-05）：產品別營收（雷射、光模組、OCS）公司不揭露，本版為 Assumed 拆分。網通 CPO 已改三情境（v0.2.2；A＝替代快、C＝慢）：CPO 快對 Lumentum 外部光源是加分，情境 A FY2028 EPS 由 31.43 升到 32.73（+4.1%）、C 由 25.45 降到 25.04（−1.6%）、B 不變；CPO 外部光源份額低／高時 FY2028 EPS −8.2%／+5.4%。"
      ]   # 重點觀察（與 README.md 相同文字）
_n0 = RD.max_row
RD.cell(_n0 + 1, 1, "v0.2.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（Val 頁 0.3 結論加「上檔空間級距」「投資建議」兩列、README 頁加重點觀察三點；數值不變，既有儲存格位置不變）。")
RD.cell(_n0 + 2, 1, f'="投資建議："&Val!B{RC_ROW}&"（依上檔空間固定規則）"').font = BOLD
for _i, _t in enumerate(OBS, 1): RD.cell(_n0 + 2 + _i, 1, f"重點觀察 {_i}：" + _t)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# ---------------- v0.2.2（2026-10-26）：網通 CPO 三情境（同字母）；In 頁在現有最後一列之後新增，不移動既有列 ----------------
_ri = IN.max_row + 2
IN.cell(_ri, 1, "網通 CPO 三情境（v0.2.2；markets/network v0.2.3：A＝替代快、B＝平平、C＝慢；情境 A 為上方 A 節，B、C 為下列，各再乘同字母情境顆數係數）").font = BOLD; _ri += 1
for _kb, _lab, _kind in [('mpl', 'AI 可插拔光模組市場', 'mpl'), ('mcp', 'AI CPO 光學市場', 'mcp')]:
    for _s in ("B", "C"):
        for _j, _y in enumerate((26, 27, 28, 29)):
            _k = f"{_kb}{_s}_{_y}"
            IN.cell(_ri, 1, _k); IN.cell(_ri, 2, f"情境 {_s}：{_lab} CY20{_y}"); IN.cell(_ri, 3, "US$B")
            IN.cell(_ri, 5, NETS[(_kind, _s)][0][_j]).font = BLUE
            IN.cell(_ri, 7, "Derived"); IN.cell(_ri, 8, NETS[(_kind, _s)][1])
            IN.cell(_ri, 9, f"網通模型 CPO 情境 {_s}（技術替代屬行業情境，公司讀同字母）" if _y == 26 else "同上")
            wb.defined_names[_k] = DefinedName(_k, attr_text=f"In!$E${_ri}"); _ri += 1
_n1 = RD.max_row
RD.cell(_n1 + 1, 1, 'v0.2.2（2026-10-26）：網通模型 v0.2.3 的 CPO 滲透改三情境（A＝替代快、B＝平平、C＝慢），本模型 A／B／C 改讀同字母的可插拔與 CPO 市場（In 頁末新增 B、C 兩組；情境 A 沿用 A 節 mpl_／mcp_，數值改為網通情境 A）；FY2027 起三情境不再相同。其他輸入不變。')
wb.save(sys.argv[1])
if os.environ.get("LITE_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK_ROW, "SENSN": SENS_N, "SENSNOTE": SENS_NOTE, "VAL": VAL,
               "IMPL": {str(k): v for k, v in IMPL.items()}, "NOTE4": NOTE4, "NOPRIOR": NOPRIOR, "RECON_END": RECON_END, "RULE": rule_rows,
               "IN": INROW},
              open(os.environ["LITE_ROWMAP"], "w"), ensure_ascii=False, indent=1)
