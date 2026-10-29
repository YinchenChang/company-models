# Astera Labs（ALAB）收支與估值模型 v0.1.1（AI 半導體 Project：網通次產業內複製——PCIe／CXL 重定時器、Scorpio 交換晶片、AEC 模組；scale-up 互連）
# 用法：python3 alab/build_alab.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年（12 月底結束）；上游市場層是日曆年，FY(t) 市場＝CY(t)（不需月數加權）。
# 金額：十億美元（US$B）；EPS：美元；口徑：公司 non-GAAP（排除股票薪酬、併購費用與非現金公允價值調整）；十年 P/E 規則值用 GAAP 換算 EPS。
# FY2025 為完整實績；FY2026 錨點＝1Q26、2Q26 實績＋3Q26 指引中點＋4Q26（季增 g_q4，Assumed）。3Q26 財報 2026-11-03 公布。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
_nwb = openpyxl.load_workbook(NET_F, data_only=True)
NET_LABS = {"scale-up 交換晶片": "su", "AI IT 資本支出（加速器 × 每顆 IT capex）": "itc"}
NETR = {}
for r in _nwb["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in NET_LABS:
        NETR[NET_LABS[lab]] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETR) == 2 and all(v is not None for k in NETR for v in NETR[k][0]), NETR.keys()
# 加速器顆數（Units 頁，日曆年 2026–2029＝B–E 欄）
UNR = {}
for r in _nwb["Units"].iter_rows(min_row=4):
    if r[0].value: UNR[r[0].value] = (r[0].row, [float(r[1 + j].value) for j in range(4)])
assert {"合計", "AMD Instinct", "AWS Trainium", "其他 ASIC（MTIA、Maia、Ascend 等）", "Google TPU"} <= set(UNR), UNR.keys()
NIN = {}
for r in _nwb["In"].iter_rows(min_row=4):
    if r[0].value: NIN[r[0].value] = (r[0].row, r[4].value)
# 非 NVIDIA、非 TPU 的 scale-up 交換晶片市場＝Σ 顆數 × 每顆 scale-up 交換晶片內容（AMD 2026 為 8 GPU 伺服器、2027 起 Helios 機架；與網通模型 TYPE2PLAT 相同）
AMDP = {0: "AMD8", 1: "AMDR", 2: "AMDR", 3: "AMDR"}
NNV = [("AMD Instinct", None), ("AWS Trainium", "TRN"), ("其他 ASIC（MTIA、Maia、Ascend 等）", "ASIC")]
MSX = []; NNU = []
for j in range(4):
    tot = 0.0; uu = 0.0
    for nm, p in NNV:
        pp = AMDP[j] if p is None else p
        tot += UNR[nm][1][j] * float(NIN[f"c_{pp}_su_sw"][1]) / 1000; uu += UNR[nm][1][j]
    MSX.append(tot); NNU.append(uu)
MSX_SRC = (f"{rel(NET_F)}｜Units!B{UNR['AMD Instinct'][0]}:E{UNR['AMD Instinct'][0]}、B{UNR['AWS Trainium'][0]}:E{UNR['AWS Trainium'][0]}、"
           f"B{UNR['其他 ASIC（MTIA、Maia、Ascend 等）'][0]}:E{UNR['其他 ASIC（MTIA、Maia、Ascend 等）'][0]} × In!E{NIN['c_AMD8_su_sw'][0]}（2026）／E{NIN['c_AMDR_su_sw'][0]}、E{NIN['c_TRN_su_sw'][0]}、E{NIN['c_ASIC_su_sw'][0]}")
UTOT = UNR["合計"]
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝SEC 8-K 新聞稿與 10-K；公司不揭露產品線營收（Aries、Taurus、Leo、Scorpio P／X）與客戶名稱"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
PR2 = "https://www.sec.gov/Archives/edgar/data/0001736297/000173629726000033/q226exhibit991.htm"
IR2 = "https://ir.asteralabs.com/news-releases/news-release-details/astera-labs-reports-second-quarter-2026-financial-results"
PR4 = "https://www.sec.gov/Archives/edgar/data/1736297/000173629726000005/q425exhibit991.htm"
K10 = "https://www.sec.gov/Archives/edgar/data/1736297/000173629726000010/alab-20251231.htm"
MB = "https://www.marketbeat.com/instant-alerts/astera-labs-q2-earnings-call-highlights-2026-08-04/"
FUT = "https://futurumgroup.com/insights/astera-labs-q2-fy-2026-earnings-pull-scorpio-switch-leadership-forward/"
FOOL = "https://www.fool.com/earnings/call-transcripts/2026/05/05/astera-labs-alab-q1-2026-earnings-transcript/"
TKB = "https://www.tikr.com/blog/astera-labs-what-a-224-bull-case-looks-like"
YH = "https://finance.yahoo.com/quote/ALAB/analysis/"
SA = "https://stockanalysis.com/stocks/alab/"
SAF = "https://stockanalysis.com/stocks/alab/forecast/"
SAR = "https://stockanalysis.com/stocks/alab/financials/ratios/"
SAB = "https://stockanalysis.com/stocks/alab/financials/balance-sheet/"
CMC = "https://companiesmarketcap.com/astera-labs/pe-ratio/"
MSC = "https://uk.marketscreener.com/quote/stock/ASTERA-LABS-INC-167508773/finances/"
TF = "https://www.tickflow.io/stock/ALAB/forecast"

sec("A. 市場（讀 markets/network 成品；日曆年 CY2026–CY2029＝會計年度；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"msx_{y}", f"非 NVIDIA scale-up 交換晶片市場 CY20{y}（AMD、Trainium、其他 ASIC；不含 TPU）", "US$B", None, MSX[j], None, "Derived", MSX_SRC if y == 26 else "同上",
        "顆數 × 每顆 scale-up 交換晶片內容（網通模型 In 頁 c_*_su_sw）；NVIDIA 平台用自家 NVLink 交換晶片、TPU 用 ICI／光學電路交換，不列入 Scorpio X 可服務市場" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"msu_{y}", f"scale-up 交換晶片市場合計 CY20{y}（含 NVIDIA NVLink；Units、Recon 用）", "US$B", None, NETR["su"][0][j], None, "Derived", NETR["su"][1], "只作對照，不進損益" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"u_{y}", f"AI 加速器顆數合計 CY20{y}", "百萬顆", None, UTOT[1][j], None, "Derived", f"{rel(NET_F)}｜Units!{'BCDE'[j]}{UTOT[0]}",
        "網通模型 Units 頁（顆數取自 markets/memory，2029 為成長率）" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"nnu_{y}", f"非 NVIDIA、非 TPU 加速器顆數 CY20{y}（AMD、Trainium、其他 ASIC）", "百萬顆", None, NNU[j], None, "Derived", MSX_SRC.split(" × ")[0] if y == 26 else "同上", "Units、Recon 用：Scorpio X 每顆內容" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 CY20{y}", "US$B", None, NETR["itc"][0][j], None, "Derived", NETR["itc"][1], "Recon 用：Astera 營收 ÷ AI IT 資本支出" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：CY2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、Credo、Arista 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：CY2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")

sec("B. Scorpio X（scale-up 交換晶片；本專案為權威：Astera 在非 NVIDIA scale-up 交換晶片的份額，規則第 13 條）")
inp("x_25", "FY2025 Scorpio X 營收", "US$B", 0.0, 0.0, 0.02, "Assumed", TKB + "：Scorpio X 在 FY2025「not meaningful」（Interested-party）", FOOL + "：Scorpio X 初期量產於 1Q26 開始出貨（公司）")
inp("x_26", "FY2026 Scorpio X 營收（未揭露）", "US$B", 0.28, 0.38, 0.50, "Assumed", "公司不揭露產品線營收：未找到；推估：Scorpio 1Q26 約占 20%（分析師於法說引述，公司未證實）、3Q26 成為最大產品線且 X 超過 P",
    MB + "：Scorpio 於 3Q26 成為最大產品線（提前一季）、X 系列 3Q26 超過 P 系列；逾 10 家客戶接洽 Scorpio X（公司）")
inp("rel_x_27", "Scorpio X 有效份額 FY2027 ÷ FY2026（相對值）", "x", 0.70, 1.00, 1.40, "Assumed", MB + "：第二家以後的客戶 2026 年底起出貨、2027 年初更多客戶量產（公司）", "競爭：Broadcom Tomahawk Ultra、NVLink Fusion、其他 UALink 交換晶片（10-K 列名競爭者含 Broadcom、Marvell）")
inp("rel_x_28", "Scorpio X 有效份額 FY2028 ÷ FY2026", "x", 0.60, 1.10, 1.70, "Assumed", FOOL + "：現有設計案每顆加速器內容逾 1,000 美元（公司）", "整合 CPO 的 Scorpio X 規劃 2028 年後（公司，Futurum 轉述）")
inp("rel_x_29", "Scorpio X 有效份額 FY2029 ÷ FY2026", "x", 0.50, 1.15, 2.00, "Assumed", "同上", "—")

sec("C. PCIe 與訊號調節（Aries 重定時器、Scorpio P PCIe 交換晶片、Taurus AEC 模組、Leo CXL；讀加速器顆數）")
inp("c_pcie", "每顆加速器 PCIe／CXL／AEC 可服務內容（只作份額顯示；不影響損益）", "US$/顆", 150, 250, 400, "Assumed", "第一手每顆內容：未找到", "損益＝FY2026 營收 × 顆數成長 × 相對值，與本列無關（校準時抵銷）")
inp("rel_p_27", "PCIe 等有效份額 × 每顆內容 FY2027 ÷ FY2026（相對值）", "x", 0.65, 0.80, 1.00, "Assumed", K10 + "：競爭者 Broadcom、Credo、Marvell、Microchip、Montage、Parade、Rambus（公司）",
    "PCIe 6／7 單價上升 vs 重定時器份額競爭、平台整合（每顆加速器重定時器數減少）；Aries FY2025 年增約 70%（TIKR，Interested-party）")
inp("rel_p_28", "PCIe 等有效份額 × 每顆內容 FY2028 ÷ FY2026", "x", 0.55, 0.75, 1.00, "Assumed", "同上；Taurus 3.2T 可服務市場 2030 年逾 40 億美元（公司，MarketBeat 轉述）", "Leo CXL 2027 年起兩家美國超大規模業者量產（公司）")
inp("rel_p_29", "PCIe 等有效份額 × 每顆內容 FY2029 ÷ FY2026", "x", 0.45, 0.70, 1.00, "Assumed", "同上", "—")
inp("cust_top", "最大終端客戶占營收（FY2025；Sens「最大客戶部分轉單」用）", "%", None, 0.70, None, "Verified", K10 + "：單一終端客戶「逾 70%」、前三大約 86%（FY2025，公司未具名）", "客戶名稱：未找到（公司未具名）")
inp("cl_loss", "最大客戶部分轉單：其 PCIe 等內容轉給其他供應商或自製的比例（Sens 用）", "%", None, 0.30, None, "Assumed", "無", "—")
for y in (27, 28, 29):
    inp(f"cl_{y}", f"最大客戶部分轉單情境：PCIe 等相對值 FY20{y}（Sens 用）", "x", None, f"=rel_p_{y}*(1-cust_top*cl_loss)", None, "Derived", "相對值 ×（1 − 最大客戶占比 × 轉單比例）", "同上" if y > 27 else "假設轉單只發生在 PCIe 等（Scorpio X 不變）")

sec("D. FY2026 錨點（1Q26、2Q26 實績＋3Q26 指引＋4Q26 假設）")
inp("q1_rev", "1Q26 營收（截至 2026-03-31）", "US$B", None, 0.308361, None, "Verified", PR2 + "：308,361 千美元（2Q26 新聞稿比較欄）", FOOL + "：1Q26 法說")
inp("q1_gp", "1Q26 non-GAAP 毛利", "US$B", None, 0.235640, None, "Verified", PR2 + "：235,640 千美元（毛利率 76.4%）", FOOL + "：76.4%")
inp("q1_opx", "1Q26 non-GAAP 營業費用（＝毛利 − 營業利益）", "US$B", None, "=0.235640-0.111729", None, "Verified", PR2 + "：non-GAAP 營業利益 111,729 千美元", "—")
inp("q1_ninc", "1Q26 利息與其他收益（non-GAAP）", "US$B", None, 0.011581, None, "Verified", PR2 + "：11,581 千美元", "—")
inp("q1_ni", "1Q26 non-GAAP 淨利", "US$B", None, 0.110069, None, "Verified", PR2 + "：110,069 千美元", "—")
inp("q1_sh", "1Q26 稀釋股數", "十億股", None, 0.181157, None, "Verified", PR2 + "：181,157 千股", "—")
inp("q2_rev", "2Q26 營收（截至 2026-06-30）", "US$B", None, 0.392400, None, "Verified", PR2 + "：392,400 千美元（年增 104%、季增 27%）", MB + "：3.924 億美元")
inp("q2_gp", "2Q26 non-GAAP 毛利", "US$B", None, 0.289259, None, "Verified", PR2 + "：289,259 千美元（毛利率 73.7%；含 Amazon 認股權證約 200bp 非現金營收減項）", FUT + "：73.7%")
inp("q2_opx", "2Q26 non-GAAP 營業費用（＝毛利 − 營業利益）", "US$B", None, "=0.289259-0.153472", None, "Verified", PR2 + "：non-GAAP 營業利益 153,472 千美元（營業利益率 39.1%）", MB + "：營業費用 1.358 億美元")
inp("q2_ninc", "2Q26 利息與其他收益（non-GAAP；扣除 150 萬美元公允價值調整）", "US$B", None, "=0.013577-0.0015", None, "Verified", PR2 + "：GAAP 13,577 千美元；non-GAAP 調整 −1,500", MB + "：利息與其他收益 1,210 萬美元")
inp("q2_ni", "2Q26 non-GAAP 淨利", "US$B", None, 0.145816, None, "Verified", PR2 + "：145,816 千美元", FUT + "：1.458 億美元")
inp("q2_sh", "2Q26 稀釋股數", "十億股", None, 0.183340, None, "Verified", PR2 + "：183,340 千股", "—")
inp("q2_eps", "2Q26 non-GAAP 稀釋 EPS", "US$", None, 0.80, None, "Verified", PR2 + "：0.80（GAAP 0.83，含所得稅利益）", FUT + "：0.80")
inp("q3_rev", "3Q26 營收指引（中點；截至 2026-09-30）", "US$B", 0.540, 0.550, 0.560, "Interested-party", PR2 + "：5.40–5.60 億美元（Scorpio X 320 通道交換晶片量產）", MB + "：同值，中點季增約 40%")
inp("q3_gm", "3Q26 non-GAAP 毛利率指引", "%", 0.715, 0.72, 0.725, "Interested-party", PR2 + "：約 72%", FUT + "：硬體比重較高的交換晶片占比上升")
inp("q3_opx", "3Q26 non-GAAP 營業費用指引（中點）", "US$B", 0.156, 0.158, 0.160, "Interested-party", PR2 + "：1.56–1.60 億美元", MB + "：同值")
inp("q3_eps", "3Q26 non-GAAP EPS 指引（中點）", "US$", 1.16, 1.185, 1.21, "Interested-party", PR2 + "：1.16–1.21", MB + "：同值；稀釋股數約 1.85 億股")
inp("q3_sh", "3Q26 稀釋股數指引", "十億股", None, 0.185, None, "Interested-party", PR2 + "：約 1.85 億股", "—")
inp("tx_26", "FY2026 下半年 non-GAAP 稅率", "%", 0.11, 0.12, 0.13, "Interested-party", PR2 + "：3Q26 non-GAAP 稅率約 12%", "1H26 推導約 11%")
inp("q_oth", "3Q26、4Q26 每季利息與其他收益", "US$B", 0.011, 0.0125, 0.014, "Assumed", "2Q26 0.0121；現金與有價證券 12.5 億美元（2026-06-30）", "—")
inp("g_q4", "4Q26 營收季增", "%", 0.05, 0.15, 0.30, "Assumed", "公司未給 4Q26 指引；Scorpio X 第二家以後客戶 2026 年底起出貨（公司）", YH + "：共識 4Q26 營收 6.557 億美元（18 家；6.09–7.38；Interested-party，不作校準）")
inp("gm_q4", "4Q26 non-GAAP 毛利率", "%", 0.70, 0.715, 0.725, "Assumed", "3Q26 指引 72%；公司長期目標 70%（公司，MarketBeat 轉述）", "—")
inp("q4_opx", "4Q26 non-GAAP 營業費用", "US$B", 0.162, 0.168, 0.176, "Assumed", "3Q26 指引中點 1.58 億美元；2Q→3Q 季增約 16%", "—")
inp("q4_sh", "4Q26 稀釋股數", "十億股", 0.1855, 0.1865, 0.188, "Assumed", "3Q26 約 1.85 億股；Amazon 認股權證最多 326 萬股（TIKR 轉述，Interested-party）", "—")

sec("E. FY2024、FY2025 實績（non-GAAP）與股東權益")
for y, rv, gp, oi, ninc, ni, sh, eps, src2 in (
        (24, 0.396290, 0.303544, 0.119594, 0.034288, 0.143329, 0.168913, 0.84, "非 GAAP 擬制稀釋股數 168,913 千股（2024-03 上市前後）；GAAP 淨損 83,421"),
        (25, 0.852525, 0.646384, 0.334406, 0.044730, 0.331015, 0.179551, 1.84, "GAAP 淨利 219,134、GAAP 稀釋 EPS 1.22")):
    inp(f"a{y}_rev_in", f"FY20{y} 營收", "US$B", None, rv, None, "Verified", PR4, MSC + "：同值（" + ("396.3" if y == 24 else "852.5") + " 百萬美元）")
    inp(f"a{y}_gp_in", f"FY20{y} non-GAAP 毛利", "US$B", None, gp, None, "Verified", PR4 + ("：non-GAAP 毛利率 76.6%" if y == 24 else "：non-GAAP 毛利率 75.8%"), "—")
    inp(f"a{y}_opx_in", f"FY20{y} non-GAAP 營業費用（＝毛利 − 營業利益）", "US$B", None, f"={gp}-{oi}", None, "Verified", PR4 + f"：non-GAAP 營業利益 {oi * 1e6:,.0f} 千美元", "—")
    inp(f"a{y}_ninc_in", f"FY20{y} 利息收入", "US$B", None, ninc, None, "Verified", PR4, "—")
    inp(f"a{y}_ni_in", f"FY20{y} non-GAAP 淨利", "US$B", None, ni, None, "Verified", PR4, src2)
    inp(f"a{y}_sh_in", f"FY20{y} non-GAAP 稀釋股數", "十億股", None, sh, None, "Verified", PR4, "—")
    inp(f"a{y}_eps_in", f"FY20{y} non-GAAP 稀釋 EPS", "US$", None, eps, None, "Verified", PR4 + ("：0.84（GAAP −0.64）" if y == 24 else "：1.84（GAAP 1.22）"), "—")
inp("a25_ni_h1", "1H25 non-GAAP 淨利（近 12 月 ROE 用）", "US$B", None, 0.137650, None, "Verified", PR2 + "：1H25 137,650 千美元", "—")
for k, nm, v, s2 in (("eq_23", "股東權益 2023-12-31", 0.15743, "上市前（2024-03 IPO）"), ("eq_24", "股東權益 2024-12-31", 0.9648, "—"),
                     ("eq_25", "股東權益 2025-12-31", 1.363634, PR2 + "：1,363,634 千美元"), ("eq_q2", "股東權益 2026-06-30", 1.725704, PR2 + "：1,725,704 千美元")):
    inp(k, nm, "US$B", None, v, None, "Verified", SAB, s2)

sec("G. 營業費用、毛利率、稅率、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.40, "=Q!B28", 0.90, "Derived", "FY2026 錨點營業費用年增 ÷ 營收年增（Q 頁）", "FY2025：69.6% ÷ 115.1% ＝ 0.60（推導）")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.01, 0.0, 0.02, "Assumed", "無", "Sens 用 +2 個百分點")
inp("gm_27", "non-GAAP 毛利率 FY2027", "%", 0.66, 0.70, 0.72, "Assumed", MB + "：長期毛利率目標 70%（公司）", "FY2024 76.6%、FY2025 75.8%、3Q26 指引 72%（交換晶片硬體占比上升）")
inp("gm_28", "non-GAAP 毛利率 FY2028", "%", 0.65, 0.695, 0.72, "Assumed", "同上", "—")
inp("gm_29", "non-GAAP 毛利率 FY2029", "%", 0.64, 0.69, 0.72, "Assumed", "同上；競爭者增加時價格下壓", "—")
inp("oth_f", "利息與其他收益（每年，FY2027 起）", "US$B", 0.03, 0.05, 0.07, "Assumed", "現金與有價證券 12.5 億美元（2026-06-30）× 約 4%", "—")
inp("tx_27", "non-GAAP 有效稅率 FY2027", "%", 0.10, 0.12, 0.15, "Assumed", PR2 + "：non-GAAP 稅率約 12%（公司）", "FY2025 推導 12.7%")
inp("tx_28", "non-GAAP 有效稅率 FY2028", "%", 0.10, 0.12, 0.15, "Assumed", "同上", "—")
inp("tx_29", "non-GAAP 有效稅率 FY2029", "%", 0.10, 0.12, 0.15, "Assumed", "同上", "—")
inp("sh_g", "股數年變動（FY2027 起）", "%", 0.005, 0.015, 0.03, "Assumed", "FY2025→1H26 稀釋股數 179.6→183.3 百萬股；1H26 股票薪酬 1.13 億美元", "買回計畫：未找到")
inp("mf", "非 NVIDIA scale-up 交換晶片市場係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 CY2027 非 NVIDIA scale-up 交換 +279%（AMD 改機架）、CY2028 +23%、CY2029 +15%", "±20%（NVLink Fusion、乙太網路 scale-up 搶占 UALink／PCIe 架構）")
inp("g2n", "GAAP EPS ÷ non-GAAP EPS（規則值換算用）", "x", 0.62, 0.70, 0.78, "Assumed", PR4 + "：FY2025 GAAP 1.22 ÷ non-GAAP 1.84 ＝ 0.66", MSC + "：FY2027 共識 GAAP 4.75 ÷ Yahoo non-GAAP 6.40 ＝ 0.74（Interested-party）")

sec("H. 共識、市場資料（只作 Recon 對帳，不作校準）")
inp("cons_q3", "共識 3Q26 營收", "US$B", None, 0.5508, None, "Interested-party", YH + "（18 家；5.457–5.570）", "—")
inp("cons_q4", "共識 4Q26 營收", "US$B", None, 0.65571, None, "Interested-party", YH + "（18 家；6.091–7.379）", "—")
inp("cons_rev26", "共識 FY2026 營收", "US$B", None, 1.91, None, "Interested-party", YH + "（21 家；1.86–1.99）", SAF + "：1.91（21 家）")
inp("cons_eps26", "共識 FY2026 non-GAAP EPS", "US$", None, 4.03, None, "Interested-party", YH + "（21 家；3.90–4.33）", SAF + "：4.03")
inp("cons_rev27", "共識 FY2027 營收", "US$B", None, 2.98, None, "Interested-party", YH + "（24 家；2.19–3.68）", MSC + "：2.976")
inp("cons_eps27", "共識 FY2027 non-GAAP EPS", "US$", None, 6.40, None, "Interested-party", YH + "（23 家；4.35–8.80）", "第二來源（non-GAAP）：未找到；MarketScreener GAAP 4.749")
inp("cons_rev28", "共識 FY2028 營收", "US$B", None, 3.835, None, "Interested-party", MSC + "（26 家）", "第二來源：未找到")
inp("cons_eps28", "共識 FY2028 non-GAAP EPS（推導＝MarketScreener GAAP 6.36 × FY2027 non-GAAP ÷ GAAP）", "US$", None, "=6.36*cons_eps27/4.749", None, "Derived",
    MSC + "：FY2028 GAAP EPS 6.36、FY2027 4.749（Interested-party）", "non-GAAP FY2028 共識：未找到（stockanalysis 需付費）")
inp("tp_street", "賣方平均目標價", "US$", None, 388.36, None, "Interested-party", SAF + "（26 家，資料 2026-10-16；190–500；中位數 425）", TF + "：370.42（24 家；頁面數字前後不一致）")
inp("tam_su", "公司估計：商用 scale-up 交換晶片市場 2030 年", "US$B", None, 20, None, "Interested-party", PR4 + "：2030 年每年 200 億美元（公司）", "—")
inp("tp_wA", "目標價情境權重 A", "%", None, 0.3, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")
inp("tp_wB", "目標價情境權重 B", "%", None, 0.5, None, "Assumed", "Andy 可改", "—")
inp("tp_wC", "目標價情境權重 C", "%", None, 0.2, None, "Assumed", "Andy 可改", "—")

sec("I. 價格與估值")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 257.16, None, "Interested-party", SA + "：257.16（市值 446.1 億美元）", YH + "：257.16")
inp("px_hi", "52 週高點", "US$", None, 499.48, None, "Interested-party", SA + "：52 週區間 97.89–499.48", "—")
inp("sh_out", "流通股數（2026-10）", "十億股", None, 0.17349, None, "Interested-party", SA + "：173.49 百萬股", K10 + "：170,214,503 股（2026-01-31，公司）")

import re
NAMES = {PR2: "公司 2Q26 新聞稿", IR2: "公司 IR 2Q26", PR4: "公司 4Q25 新聞稿", K10: "公司 FY2025 10-K", MB: "MarketBeat 法說摘要", FUT: "Futurum", FOOL: "Motley Fool 法說逐字稿",
         TKB: "TIKR", YH: "Yahoo Finance", SA: "StockAnalysis", SAF: "StockAnalysis 預估", SAR: "StockAnalysis 比率", SAB: "StockAnalysis 資產負債表", CMC: "companiesmarketcap",
         MSC: "MarketScreener", TF: "Tickflow"}
def strip_url(t):   # 含「未找到」的列會被 HTML 掃描列出：網址改寫為來源名稱（網址見 Sources）
    if not isinstance(t, str): return t
    for u, nm in NAMES.items(): t = t.replace(u, nm + "（網址見 Sources）")
    return re.sub(r"https?://[^\s：（）]+", "（網址見 Sources）", t)
sec("K. 目標價（DCF＋Comparable）與收斂法參數（黃底，Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；各市場同一值）", "DCF 與 Sens")
inp("tp_g", "長期成長率（DCF 永續；收斂法情境 B 的 FY2032 營收年增）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "—")
inp("tp_roe", "ROE（與 EPS 同口徑：non-GAAP 淨利 ÷ 平均股東權益，FY2024、FY2025、近 12 月〔至 2Q26〕三期平均；DCF 再投資率＝tp_g ÷ tp_roe）", "%", None,
    "=AVERAGE(a24_ni_in/AVERAGE(eq_23,eq_24),a25_ni_in/AVERAGE(eq_24,eq_25),(a25_ni_in-a25_ni_h1+q1_ni+q2_ni)/AVERAGE(eq_25,eq_q2))", None, "Derived",
    "公司新聞稿 non-GAAP 淨利 ÷ StockAnalysis 年底股東權益（FY2024 權益因 2024-03 上市大增）", "參考（不採）：GAAP ROE stockanalysis FY2025 18.82%、近 12 月 25.83%")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（工作單 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "收斂法：情境 A 於 FY2032 的營收年增（＝tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法；FY2030 起由 FY2029 年增線性收斂", "FY2030 起為收斂法，非驅動因子推導")
inp("fade_gB", "收斂法：情境 B 於 FY2032 的營收年增（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "收斂法：情境 C 於 FY2032 的營收年增（＝tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "收斂法：情境 A 的中段淨利率（＝情境 A FY2029 淨利率）", "%", None, None, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "收斂法：情境 B 的中段淨利率（＝FY2029 與十年平均的中點；十年平均以 FY2029 × 0.8 代替）", "%", None, None, None, "Assumed", "十年平均淨利率（2016–2025，與本模型同口徑）：未找到（2024-03 上市，2023 年以前虧損），依工作單以 FY2029 × 0.8 代替；本列＝（FY2029＋FY2029 × 0.8）÷ 2", "Andy 可改")
inp("fade_mC", "收斂法：情境 C 的中段淨利率（＝十年平均；以情境 C FY2029 × 0.8 代替）", "%", None, None, None, "Assumed", "十年平均淨利率：未找到，依工作單以情境 C FY2029 淨利率 × 0.8 代替", "Andy 可改")
inp("cons_tp", "市場共識目標價（賣方平均）", "US$", None, "=tp_street", None, "Interested-party", "賣方平均目標價（In 頁 tp_street：stockanalysis 26 家 388.36，資料 2026-10-16）", "Tickflow 24 家 370.42")
inp("cons_feps", "共識 12 個月後前瞻 EPS（＝2/12 × FY2027＋10/12 × FY2028 共識）", "US$", None, "=2/12*cons_eps27+10/12*cons_eps28", None, "Derived", "Yahoo FY2027 6.40；FY2028 為推導值（In 頁 cons_eps28）", "共識只放 Recon 與本列，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    if any("未找到" in str(x) for x in (item, s1, s2)): s1, s2 = strip_url(s1), strip_url(s2)
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 64, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億美元；EPS 為美元；股數為十億股；non-GAAP）。藍字＝輸入；黑字＝公式。會計年度＝日曆年"
Q["A2"] = "FY2026 錨點＝1Q26、2Q26 實績＋3Q26 指引中點（營收 5.50 億美元、毛利率 72%、營業費用 1.58 億美元）＋4Q26（季增 g_q4、毛利率 gm_q4，Assumed）。產品線拆分（Scorpio X）為假設，PCIe 與訊號調節為殘差。3Q26 財報（2026-11-03）公布後以實績取代 3Q"
hdr(Q, 3, ["項目", "FY2024A", "FY2025A", "1Q26（實績）", "2Q26（實績）", "3Q26（指引）", "4Q26（假設）", "FY2026 錨點", "標記", "說明"])
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", note=""):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
    Q.cell(r, 9, tag); Q.cell(r, 10, note)
QC = "BCDEFGH"
qrow(4, "營收", {"B": "=a24_rev_in", "C": "=a25_rev_in", "D": "=q1_rev", "E": "=q2_rev", "F": "=q3_rev", "G": "=F4*(1+g_q4)", "H": "=SUM(D4:G4)"}, tag="Verified／Interested-party（指引）／Assumed", note="4Q26＝3Q26 ×（1＋g_q4）")
qrow(5, "毛利（non-GAAP）", {"B": "=a24_gp_in", "C": "=a25_gp_in", "D": "=q1_gp", "E": "=q2_gp", "F": "=F4*q3_gm", "G": "=G4*gm_q4", "H": "=SUM(D5:G5)"}, tag="Verified／Derived")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in QC}, fmt="0.0%", tag="Derived")
qrow(7, "營業費用（non-GAAP）", {"B": "=a24_opx_in", "C": "=a25_opx_in", "D": "=q1_opx", "E": "=q2_opx", "F": "=q3_opx", "G": "=q4_opx", "H": "=SUM(D7:G7)"}, tag="Verified／Interested-party（指引）／Assumed")
qrow(8, "營業費用率", {c: f"={c}7/{c}4" for c in QC}, fmt="0.0%", tag="Derived")
qrow(9, "營業利益（non-GAAP）", {c: f"={c}5-{c}7" for c in QC}, tag="Derived")
qrow(10, "利息與其他收益", {"B": "=a24_ninc_in", "C": "=a25_ninc_in", "D": "=q1_ninc", "E": "=q2_ninc", "F": "=q_oth", "G": "=q_oth", "H": "=SUM(D10:G10)"}, tag="Verified／Assumed")
qrow(11, "稅前利益", {c: f"={c}9+{c}10" for c in QC}, tag="Derived")
qrow(12, "所得稅（non-GAAP）", {"B": "=B11-a24_ni_in", "C": "=C11-a25_ni_in", "D": "=D11-q1_ni", "E": "=E11-q2_ni", "F": "=F11*tx_26", "G": "=G11*tx_26", "H": "=SUM(D12:G12)"}, tag="Derived", note="實績＝稅前 − 淨利（推導）")
qrow(13, "有效稅率", {c: f"={c}12/{c}11" for c in QC}, fmt="0.0%", tag="Derived")
qrow(14, "淨利（non-GAAP）", {"B": "=a24_ni_in", "C": "=a25_ni_in", "D": "=q1_ni", "E": "=q2_ni", "F": "=F11-F12", "G": "=G11-G12", "H": "=SUM(D14:G14)"}, tag="Verified／Derived")
qrow(15, "淨利率", {c: f"={c}14/{c}4" for c in QC}, fmt="0.0%", tag="Derived")
qrow(16, "稀釋股數（non-GAAP）", {"B": "=a24_sh_in", "C": "=a25_sh_in", "D": "=q1_sh", "E": "=q2_sh", "F": "=q3_sh", "G": "=q4_sh", "H": "=AVERAGE(D16:G16)"}, tag="Verified／Assumed", note="全年＝四季平均")
qrow(17, "EPS（non-GAAP）", {c: f"={c}14/{c}16" for c in QC}, fmt="0.00", tag="Derived", note="公司揭露：FY2024 0.84、FY2025 1.84、1Q26 0.61、2Q26 0.80；3Q26 指引 1.16–1.21")
qrow(18, "Scorpio X 營收", {"C": "=x_25", "H": "=x_26"}, tag="Assumed")
qrow(19, "PCIe 與訊號調節（Aries、Scorpio P、Taurus、Leo）＝營收 − Scorpio X（殘差）", {"C": "=C4-C18", "H": "=H4-H18"}, tag="Derived")
qrow(20, "Scorpio X ÷ 營收", {"C": "=C18/C4", "H": "=H18/H4"}, fmt="0.0%", tag="Derived", note="公司：Scorpio 3Q26 成為最大產品線、X 超過 P")
Q.cell(23, 1, "FY2026 錨點的檢查（參考；校準未使用共識）").font = BOLD
Q.cell(24, 1, "FY2026 營收 ÷ 共識 − 1"); Q["B24"] = "=H4/cons_rev26-1"
Q.cell(25, 1, "FY2026 EPS ÷ 共識 − 1"); Q["B25"] = "=H17/cons_eps26-1"
Q.cell(26, 1, "FY2026 營收年增"); Q["B26"] = "=H4/C4-1"
Q.cell(27, 1, "FY2026 營業費用年增"); Q["B27"] = "=H7/C7-1"
Q.cell(28, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B28"] = "=B27/B26"
Q.cell(29, 1, "2Q26 EPS（模型）vs 公司揭露 0.80"); Q["B29"] = "=E17/q2_eps-1"
Q.cell(30, 1, "3Q26 EPS（模型）vs 指引中點 1.185"); Q["B30"] = "=F17/q3_eps-1"
Q.cell(31, 1, "4Q26 營收 ÷ 共識 − 1"); Q["B31"] = "=G4/cons_q4-1"
for rr in range(24, 32): Q[f"B{rr}"].number_format = "0.0%" if rr != 28 else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("ninc", 10), ("tx", 13), ("ni", 14), ("sh", 16), ("eps", 17)):
    for y, col in ((24, "B"), (25, "C"), (26, "H")):
        defname(f"a{y}_{nm}", f"Q!${col}${rr}")
for nm, rr in (("x", 18), ("p", 19)):
    for y, col in ((25, "C"), (26, "H")):
        defname(f"a{y}_{nm}", f"Q!${col}${rr}")
Q.column_dimensions["A"].width = 56
for c in "BCDEFGH": Q.column_dimensions[c].width = 14
Q.column_dimensions["I"].width = 30; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029); F_ = (2027, 2028, 2029)
umx = lambda cy, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[cy]
def mkt(base, y, s, mf):
    """FY(y) 市場＝CY(y) × 顆數係數；FY2027 起乘市場係數 mf"""
    f = f"{base}_{y % 100}" + ("" if umx(y, s) == "1" else f"*{umx(y, s)}")
    return f"={f}" + ("" if y == 2026 or mf == "1" else f"*{mf}")
DEFP = dict(scn="A", relx="rel_x", relp="rel_p", gm="gm", tx="tx", mf="1", opxa="0", phi="phi")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.000", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. Scorpio X（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029；市場為日曆年＝會計年度）")
    row("msx", "非 NVIDIA scale-up 交換晶片市場 × 情境顆數係數", "US$B", {c_(y): mkt("msx", y, s, P["mf"]) for y in BU}, fmt="#,##0.00",
        note="AMD、Trainium、其他 ASIC 顆數 × 每顆 scale-up 交換晶片內容（網通模型）；CY2028 起乘顆數係數")
    row("shx0", "Astera Scorpio X 有效份額（FY2026 校準）＝ Scorpio X 營收 ÷ 市場", "%", {"D": f"=a26_x/D{R['msx']}"}, fmt="0.00%", note="Derived")
    row("relx", "Scorpio X 有效份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['relx']}_{y % 100}" for y in F_}}, fmt="0.00")
    row("shx", "Scorpio X 有效份額", "%", {c_(y): f"=$D${R['shx0']}*{c_(y)}{R['relx']}" for y in BU}, fmt="0.00%")
    row("x", "Scorpio X 營收", "US$B", {"C": "=a25_x", "D": "=a26_x", **{c_(y): f"={c_(y)}{R['msx']}*{c_(y)}{R['shx']}" for y in F_}})
    sec_("B. PCIe 與訊號調節（Aries、Scorpio P、Taurus、Leo）")
    row("u", "AI 加速器顆數 × 情境顆數係數", "百萬顆", {c_(y): mkt("u", y, s, "1") for y in BU}, fmt="#,##0.00")
    row("mp", "PCIe／CXL／AEC 可服務市場＝顆數 × 每顆內容（c_pcie）", "US$B", {c_(y): f"={c_(y)}{R['u']}*c_pcie/1000" for y in BU}, fmt="#,##0.00", note="c_pcie 只影響份額顯示，校準時抵銷")
    row("shp0", "Astera PCIe 等有效份額（FY2026 校準）＝ 營收 ÷ 可服務市場", "%", {"D": f"=a26_p/D{R['mp']}"}, fmt="0.00%")
    row("relp", "PCIe 等有效份額 × 每顆內容相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['relp']}_{y % 100}" for y in F_}}, fmt="0.00")
    row("p", "PCIe 與訊號調節營收", "US$B", {"C": "=a25_p", "D": "=a26_p", **{c_(y): f"={c_(y)}{R['mp']}*$D${R['shp0']}*{c_(y)}{R['relp']}" for y in F_}})
    sec_("C. 合計")
    row("rev", "營收", "US$B", {"C": "=a25_rev", **{c_(y): f"={c_(y)}{R['x']}+{c_(y)}{R['p']}" for y in BU}})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("xs", "Scorpio X ÷ 營收", "%", {c_(y): f"={c_(y)}{R['x']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    sec_("D. 損益（non-GAAP）")
    row("gm", "毛利率", "%", {"C": "=a25_gp/a25_rev", "D": "=a26_gp/a26_rev", **{c_(y): f"={P['gm']}_{y % 100}" for y in F_}}, fmt="0.0%",
        note="FY2027 起為輸入路徑（公司長期目標 70%；交換晶片硬體占比上升）")
    row("gp", "毛利", "US$B", {c_(y): f"={c_(y)}{R['rev']}*{c_(y)}{R['gm']}" for y in FYS})
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,{P['phi']}*{c_(y)}{R['g']}))" for y in F_}})
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in F_}})
    row("oi", "營業利益", "US$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("ninc", "利息與其他收益", "US$B", {"C": "=a25_ninc", "D": "=a26_ninc", **{c_(y): "=oth_f" for y in F_}})
    row("tr", "有效稅率", "%", {"C": "=a25_tx", "D": "=a26_tx", **{c_(y): f"={P['tx']}_{y % 100}" for y in F_}}, fmt="0.0%", note="公司 non-GAAP 稅率約 12%")
    row("ni", "淨利", "US$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}+{c_(y)}{R['ninc']})*(1-{c_(y)}{R['tr']})" for y in F_}})
    row("nim", "淨利率", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("sh_", "稀釋股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in F_}}, fmt="0.0000")
    row("eps", "EPS（non-GAAP）", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS}, fmt="0.00")
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Astera Labs 損益（US$B；non-GAAP）。FY2025A 實績、FY2026 錨點（1Q26、2Q26 實績＋3Q26 指引＋4Q26 假設），FY2027 起為模型。會計年度＝日曆年。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數 CY2028 起），差異在 FY2028 起"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：順利（顆數係數 1.0）", "B": "情境 B：平平（CY2028 顆數 ×0.8、CY2029 ×0.85）", "C": "情境 C：不如預期（CY2028–29 顆數 ×0.6）"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.00"), ("g", "  營收年增", "%", "0.0%"), ("x", "  Scorpio X 營收", "US$B", "#,##0.00"), ("p", "  PCIe 與訊號調節營收", "US$B", "#,##0.00"),
                                ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"), ("nim", "淨利率", "%", "0.0%"),
                                ("ni", "淨利", "US$B", "#,##0.00"), ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收、Scorpio X 營收、淨利 − 錨點（校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：Scorpio X 營收 − 市場 × 有效份額（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['x']}-D{PB['A']['msx']}*D{PB['A']['shx']}"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：（營業利益＋其他收益）×（1 − 稅率）− 淨利錨點"); PL[f"D{sr}"] = f"=(D{PB['A']['oi']}+D{PB['A']['ninc']})*(1-D{PB['A']['tr']})-a26_ni"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：市場與份額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — Astera Labs 在非 NVIDIA scale-up 交換晶片（Scorpio X）與 PCIe／CXL／AEC 的份額（本專案權威，規則第 13 條）。市場取自 markets/network（加速器顆數 × 每顆內容；日曆年＝會計年度）"
A_ = PB["A"]
uc = {2026: "C", 2027: "D", 2028: "E", 2029: "F"}
hdr(UN, 3, ["項目", "單位", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("msu", "scale-up 交換晶片市場合計（含 NVIDIA NVLink；情境 A）", "US$B", lambda y: f"=msu_{y % 100}", "#,##0.0", "markets/network Mkt 頁「scale-up 交換晶片」"),
        ("msx", "  其中非 NVIDIA、非 TPU（Scorpio X 可服務市場）", "US$B", lambda y: f"=PL!{FC[y]}{A_['msx']}", "#,##0.0", "AMD、Trainium、其他 ASIC"),
        ("x", "Astera Scorpio X 營收（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['x']}", "#,##0.00", "FY2026 為假設（公司未揭露）"),
        ("shx", "Scorpio X 有效份額（÷ 非 NVIDIA scale-up 交換晶片）", "%", lambda y: f"=PL!{FC[y]}{A_['shx']}", "0.0%", "FY2026 校準"),
        ("shs", "Scorpio X ÷ scale-up 交換晶片合計", "%", lambda y: f"=PL!{FC[y]}{A_['x']}/{uc[y]}4", "0.0%", "含 NVIDIA NVLink 的分母"),
        ("cpx", "Scorpio X 每顆非 NVIDIA 加速器營收（美元／顆）", "US$/顆", lambda y: f"=PL!{FC[y]}{A_['x']}*1000/nnu_{y % 100}", "#,##0", "公司：現有設計案每顆加速器內容逾 1,000 美元（採用該平台者）"),
        ("u", "AI 加速器顆數（情境 A）", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['u']}", "#,##0.0", "markets/network Units 頁"),
        ("p", "Astera PCIe 與訊號調節營收（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['p']}", "#,##0.00", "FY2026 為錨點殘差"),
        ("cpp", "PCIe 等每顆加速器營收（美元／顆）", "US$/顆", lambda y: f"=PL!{FC[y]}{A_['p']}*1000/PL!{FC[y]}{A_['u']}", "#,##0", "＝c_pcie × 有效份額"),
        ("shp", "PCIe 等有效份額（÷ 顆數 × c_pcie）", "%", lambda y: f"=PL!{FC[y]}{A_['p']}/PL!{FC[y]}{A_['mp']}", "0.0%", "c_pcie 250 美元／顆（Assumed）"),
        ("itr", "Astera 營收 ÷ AI IT 資本支出", "%", lambda y: f"=PL!{FC[y]}{A_['rev']}/itc_{y % 100}", "0.00%", "markets/network 的 AI IT 資本支出")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{uc[y]}{ur}"] = f(y); UN[f"{uc[y]}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
assert UR["msu"] == 4, UR
ur += 1
UN.cell(ur, 1, "讀法：Scorpio X 的 FY2026 營收為假設（公司未揭露產品線），PCIe 等為殘差；兩者的份額相對值是 Assumed，是最需要每季檢查的輸入。每顆加速器營收 ＝ Astera 營收 ÷ 顆數，可直接與公司「每顆逾 1,000 美元」的說法對照。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 70

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元）。目標價＝DCF 與 Comparable 加權（第 0 節）；舊法規則值情境加權保留參考；另列「現價 ÷ FY2027、FY2028 EPS（模型，non-GAAP）」、市場隱含 EPS 與十年 P/E 規則值。Astera Labs 2024-03 上市，年底 P/E 只有 2024–2025，且 2024 為虧損（負值，不用）：規則值只有 2025 年一個觀測值（兩個來源）"
V["A2"] = "循環位置：Astera Labs 不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 Arista、Credo）。營收到 FY2029 仍在增加時，見頂年未知，判為中段。規則值基準＝FY2027 GAAP 換算 EPS（non-GAAP × g2n），與年底 P/E 的計算口徑一致"
VAL = {}
V.cell(14, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 15, ["項目", "A", "B", "C", "說明"])
V.cell(16, 1, "現價（2026-10-23 收盤）"); V["B16"] = "=px"; V["B16"].number_format = "0.00"
V.cell(17, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(18, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(19, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(20, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((17, 2026), (18, 2027), (19, 2028), (20, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(21, 1, "現價 ÷ 共識 FY2026 EPS"); V["B21"] = "=px/cons_eps26"; V["B21"].number_format = "0.0"
V.cell(22, 1, "現價 ÷ 共識 FY2027 EPS"); V["B22"] = "=px/cons_eps27"; V["B22"].number_format = "0.0"
V.cell(23, 1, "現價 ÷ 52 週高點 − 1"); V["B23"] = "=px/px_hi-1"; V["B23"].number_format = "0%"
V.cell(24, 1, "現價 ÷ 賣方平均目標價 − 1"); V["B24"] = "=px/tp_street-1"; V["B24"].number_format = "0%"
V["E18"] = "FY2027 三情境相同（顆數係數自 CY2028 起）"; V["E19"] = "B、C 的差異自 FY2028 起"
V.cell(26, 1, "2. 循環位置（本模型營收最高年）").font = BOLD
hdr(V, 27, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 28):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(32, 1, "3. 規則值隱含股價（美元；約 2027 年底，未折現）＝ FY2027 GAAP 換算 EPS × 規則倍數").font = BOLD
hdr(V, 33, ["情境", "FY2027 EPS", "FY2028 EPS", "FY2027 GAAP 換算 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
for j, s in enumerate(SC):
    i = 34 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=PL!F{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 4, f"=B{i}*g2n").number_format = "0.00"
    V.cell(i, 5, f"=D{28 + j}").number_format = "0.0"
    V.cell(i, 6, f"=D{i}*E{i}").number_format = "#,##0"
    V.cell(i, 7, f"=F{i}/px-1").number_format = "0%"
    V.cell(i, 8, f"=px/B{i}").number_format = "0.0"
    for c in range(1, 9): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(37, 1, "FY2027＝2027 年底的近 12 月（會計年度＝日曆年）；年底 P/E 資料商以近 12 月 GAAP EPS 計算，故以 g2n 換算")
V.cell(39, 1, "4. 核心問題：市場隱含的 FY2028 EPS（＝現價 ÷ 假設倍數；non-GAAP）與模型的距離").font = BOLD
hdr(V, 40, ["假設倍數（÷ FY2028 EPS）", "市場隱含 FY2028 EPS", "模型 FY2028 EPS（A）", "隱含 ÷ 模型 A − 1", "模型 FY2028 EPS（C）", "隱含 ÷ 模型 C − 1"])
vr = 41; IMPL = {}
for m in (20, 25, 35, 45):
    V.cell(vr, 1, m).number_format = "0.0"; V.cell(vr, 1).font = BLUE
    V.cell(vr, 2, f"=px/A{vr}").number_format = "0.00"
    V.cell(vr, 3, f"=PL!F{SUMR[('A', 'eps')]}").number_format = "0.00"
    V.cell(vr, 4, f"=B{vr}/C{vr}-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[('C', 'eps')]}").number_format = "0.00"
    V.cell(vr, 6, f"=B{vr}/E{vr}-1").number_format = "0%"
    IMPL[m] = vr; vr += 1
V.cell(vr, 1, "讀法：以 25 倍（同業前瞻 P/E 中位數附近）計，現價隱含的 FY2028 EPS 高於模型 A 約三成，需要 Scorpio X 份額上升（Sens「Scorpio X 份額 高」）才達得到；以 35 倍計約等於模型 A；以 45 倍（Astera 現行前瞻 P/E 附近）計略低於模型 B。現價隱含的是情境 A（顆數不下修）且 Scorpio X 份額維持，並給成長溢價。"); READ_ROW = vr
vr += 1
V.cell(vr, 1, "本版未做先驗，無 Andy 合理 P/E 對照。"); NOPRIOR_ROW = vr
assert vr < 50, vr
PE_CMC = [None] * 8 + [None, 127.0]     # companiesmarketcap 年底 P/E；2024 −77.0（虧損，不用）；2023 年以前未上市
PE_SA = [None] * 8 + [None, 128.23]     # StockAnalysis 會計年度末（＝12 月底）P/E；2024 無（虧損）
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；只有 2025 年一個正值；2024 −77.0x 為虧損，不用）"},
    {"label": "StockAnalysis（會計年度末 P/E）", "vals": PE_SA, "tag": "Interested-party（資料商；會計年度＝日曆年，日期口徑與年底相同；2024 虧損無值）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 GAAP 換算 EPS（A）", "basis": f"=D{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 GAAP 換算 EPS（A）", "basis": f"=D{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 GAAP 換算 EPS（A）", "basis": f"=D{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝年底 P/E 依循環位置取分位。Astera Labs 只有 2025 年一個正值（約 127–128x，營收年增 115% 的高成長期），三個分位相同；歷史倍數為 GAAP 落後 P/E，規則值偏高的方向明確。前瞻 P/E 歷史待 Bloomberg。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
# 舊法：規則值情境加權（只作參考；建置後與其他內容一起下移）
V.cell(4, 1, "舊法：規則值情境加權（參考）").font = BOLD
hdr(V, 5, ["項目", "數值", "", "", "說明"])
V.cell(6, 1, "規則值情境加權價（參考，非目標價）"); V["B6"] = f"=tp_wA*F{VAL['A']}+tp_wB*F{VAL['B']}+tp_wC*F{VAL['C']}"; V["B6"].number_format = "#,##0"
V["E6"] = "權重 × 各情境隱含價（第 3 節規則值）"
V.cell(7, 1, "規則值情境加權價 ÷ 現價 − 1（參考）"); V["B7"] = "=B6/px-1"; V["B7"].number_format = "0.0%"; V["E7"] = "現價＝In 頁 px（2026-10-23 收盤）"
for rr, s, lab in ((8, "A", "情境 A 隱含價（順利）"), (9, "B", "情境 B 隱含價（平平）"), (10, "C", "情境 C 隱含價（不如預期）")):
    V.cell(rr, 1, lab); V[f"B{rr}"] = f"=F{VAL[s]}"; V[f"B{rr}"].number_format = "#,##0"
    V[f"E{rr}"] = f'="FY2027 GAAP 換算 EPS "&TEXT(D{VAL[s]},"0.00")&" × "&TEXT(E{VAL[s]},"0.0")&" 倍（第 3 節）"'
V.cell(11, 1, "情境權重 A／B／C"); V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0%"
V["E11"] = '="合計 "&TEXT(B11+C11+D11,"0%")&"；In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"'
V.cell(12, 1, "舊法說明")
V["B12"] = ("P/E：各情境 FY2027 GAAP 換算 EPS（non-GAAP EPS × g2n 0.70）× 十年年底 GAAP P/E 依循環位置取分位（非設備股：以本模型營收最高年判定，三情境皆為中段＝中位數）。"
            "Astera Labs 2024-03 上市，十年倍數只有 2025 年一個正值（約 127–128x，FY2025 營收年增 115% 的高成長期；2024 年虧損為負值不用），偏誤方向為明顯高估；隱含價約為 2027 年底價格、未折現。")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、Scorpio X 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 Scorpio X 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 Scorpio X 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("Scorpio X 份額 低", "Scorpio X 有效份額相對值 FY2027 0.70、FY2028 0.60、FY2029 0.50（Broadcom、NVLink Fusion、其他 UALink 交換晶片搶份額）", dict(relx="XLO")),
         ("Scorpio X 份額 高", "FY2027 1.40、FY2028 1.70、FY2029 2.00（逾 10 家客戶陸續量產）", dict(relx="XHI")),
         ("PCIe 等份額 低", "PCIe 等相對值 FY2027 0.65、FY2028 0.55、FY2029 0.45（重定時器價格戰、平台整合）", dict(relp="PLO")),
         ("PCIe 等份額 高", "相對值 1.00（PCIe 7 單價上升抵銷份額競爭）", dict(relp="PHI")),
         ("最大客戶部分轉單", "最大終端客戶（逾 70%）的 PCIe 等內容 30% 轉給他人或自製", dict(relp="cl")),
         ("毛利率 低", "毛利率 FY2027 66%、FY2028 65%、FY2029 64%", dict(gm="GMLO")), ("毛利率 高", "毛利率 72%", dict(gm="GMHI")),
         ("稅率 15%", "FY2027 起 non-GAAP 稅率 15%", dict(tx="TXHI")), ("稅率 10%", "FY2027 起 10%", dict(tx="TXLO")),
         ("非 NVIDIA scale-up 交換市場 ×0.8", "FY2027 起 ×0.8（NVLink Fusion、乙太網路 scale-up 搶占）", dict(mf="mf_lo")), ("非 NVIDIA scale-up 交換市場 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("營業費用彈性 0.9", "營業費用隨營收成長較多", dict(phi="phi_hi")), ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi"))]
for nm, base, suf in (("XLO", "rel_x", "_lo"), ("XHI", "rel_x", "_hi"), ("PLO", "rel_p", "_lo"), ("PHI", "rel_p", "_hi"), ("GMLO", "gm", "_lo"), ("GMHI", "gm", "_hi"), ("TXHI", "tx", "_hi"), ("TXLO", "tx", "_lo")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}_{y}"] = DefinedName(f"{nm}_{y}", attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.00"), ("D", "x", 2027, "#,##0.00"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.00"), ("I", "x", 2028, "#,##0.00"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2026 為錨點、各案例相同；顆數情境自 CY2028 起，FY2027 不受影響。Scorpio X 份額高／低是 FY2027–FY2028 EPS 最大的單一變數，其次是 PCIe 等份額與毛利率；顆數情境 C 主要打在 FY2028 起。稅率影響較小。")
SENS_READ = nn + 3
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 60
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識、公司說法、市場層。只作事後對照，不為對上數字而調參數。共識 FY2027 EPS 區間 4.35–8.80"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與指引（校準未使用共識）").font = BOLD; rr_ += 1
rec("3Q26 營收（指引中點）vs 共識", "=cons_q3", lambda s: "=Q!F4", "Yahoo（18 家）", "指引 5.40–5.60", fmt="#,##0.000")
rec("4Q26 營收（模型假設）vs 共識", "=cons_q4", lambda s: "=Q!G4", "Yahoo（18 家；6.09–7.38）", "季增 15%（Assumed）", fmt="#,##0.000")
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "Yahoo（21 家；1.86–1.99）", "")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "Yahoo（21 家；3.90–4.33）", "")
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "Yahoo（24 家；2.19–3.68）", "")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "Yahoo（23 家；4.35–8.80）", "")
rec("FY2028 營收 vs 共識", "=cons_rev28", pl("rev", 2028), "MarketScreener（26 家）", "")
rec("FY2028 EPS vs 共識（推導）", "=cons_eps28", pl("eps", 2028), "MarketScreener GAAP × non-GAAP／GAAP 比", "共識 non-GAAP FY2028：未找到")
rec("FY2027 毛利率 vs 公司長期目標 70%", 0.70, pl("gm", 2027), "公司 2Q26 法說", "", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "B. 份額與市場（機制檢查）").font = BOLD; rr_ += 1
rec("FY2028 Scorpio X 每顆非 NVIDIA 加速器營收（美元／顆）vs 公司「逾 1,000 美元」", 1000, lambda s: f"=PL!F{PB[s]['x']}*1000/nnu_28", "公司 1Q26 法說", "公司說法為採用平台的每顆內容；本列分母為全部非 NVIDIA、非 TPU 加速器", fmt="#,##0")
rec("FY2029 非 NVIDIA scale-up 交換晶片市場（網通模型）vs 公司 2030 年商用 scale-up 交換 200 億美元", "=tam_su", lambda s: f"=PL!G{PB[s]['msx']}", "公司 4Q25 新聞稿（Interested-party）", "口徑不同：公司數字為 2030 年、含所有商用 scale-up 交換", fmt="#,##0.0")
rec("FY2027 Scorpio X 有效份額（÷ 非 NVIDIA scale-up 交換）", None, pl("shx", 2027), "—", "", fmt="0.0%")
rec("FY2028 Astera 營收 ÷ AI IT 資本支出", None, lambda s: f"=PL!F{PB[s]['rev']}/itc_28", "markets/network", "每 1 美元 AI IT 資本支出分到的 Astera 營收", fmt="0.00%")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("2Q26 EPS（模型）vs 公司揭露 0.80", "=q2_eps", lambda s: "=Q!E17", "公司", "淨利 ÷ 稀釋股數", fmt="0.00")
rec("3Q26 EPS（模型）vs 指引中點", "=q3_eps", lambda s: "=Q!F17", "公司指引 1.16–1.21", "", fmt="0.00")
rec("FY2025 EPS（模型）vs 公司揭露 1.84", "=a25_eps_in", pl("eps", 2025), "公司", "", fmt="0.00")
rec("FY2026 營業費用彈性（Derived）", None, lambda s: "=Q!B28", "Q 頁", "", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 Astera 營收 ÷ 部署 GW 計算每 GW 營收"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.cell(rr_, 1, "3Q26 財報（2026-11-03）：公布後以實績取代 Q 頁 3Q 指引，並依 4Q26 指引更新 g_q4"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_END = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝SEC 8-K 新聞稿與 10-K）"; SO["A1"].font = BOLD
srcs = [("公司：2Q26 實績、3Q26 指引（8-K 附件 99.1；含 1Q26、1H 比較欄與資產負債表）", PR2),
        ("公司：2Q26 新聞稿（IR 網站）", IR2),
        ("公司：4Q25 與 FY2025 實績、1Q26 指引、2030 年商用 scale-up 交換市場 200 億美元（8-K 附件 99.1）", PR4),
        ("公司：FY2025 10-K（客戶集中、競爭者、股數）", K10),
        ("二手：2Q26 法說摘要（Scorpio 3Q26 成為最大產品線、X 超過 P、逾 10 家客戶、長期毛利率 70%、Taurus 2030 年逾 40 億美元）", MB),
        ("二手：2Q26 分析（Futurum；毛利率下滑原因、客戶集中）", FUT),
        ("二手：1Q26 法說逐字稿（每顆加速器內容逾 1,000 美元、Scorpio 1Q26 約 20% 為分析師引述）", FOOL),
        ("二手：TIKR（2026-03；Scorpio 逾 FY2025 營收 15%、Amazon 認股權證）", TKB),
        ("共識：Yahoo Finance 分析師預估（2026-10-23 頁面）", YH),
        ("共識：StockAnalysis 預估與目標價（資料 2026-10-16）", SAF),
        ("共識：MarketScreener 年度預估（FY2028 營收、GAAP EPS）", MSC),
        ("共識：Tickflow（目標價第二來源；頁面數字前後不一致）", TF),
        ("價格：StockAnalysis（2026-10-23 收盤、52 週區間、股數、前瞻 P/E）", SA),
        ("資料商：companiesmarketcap 年底 P/E", CMC),
        ("資料商：StockAnalysis 會計年度末 P/E 與 ROE", SAR),
        ("資料商：StockAnalysis 資產負債表（股東權益）", SAB),
        ("同業前瞻 P/E：StockAnalysis statistics 頁（CRDO、MRVL、AVGO、RMBS、MTSI；2026-10-23）", "https://stockanalysis.com/stocks/crdo/statistics/"),
        ("市場層：markets/network、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：3Q26 財報；產品線營收拆分（Aries、Taurus、Leo、Scorpio P／X）；客戶名稱第一手出處；每顆加速器 PCIe 內容第一手來源；FY2028 non-GAAP 共識；前瞻 P/E 歷史（Bloomberg）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Astera Labs（ALAB）收支與估值模型 v0.1.1（AI 半導體 Project：網通次產業內複製——PCIe／CXL 重定時器、Scorpio 交換晶片、AEC 模組；scale-up 互連）　2026-10-25", True),
 ("版本：v0.1（2026-10-25 初版；三情境延伸到 FY2032、目標價＝DCF 與 Comparable 加權）。Astera Labs 在非 NVIDIA scale-up 交換晶片（Scorpio X）與 PCIe／CXL／AEC 的份額在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）。", False),
 ("模型命題：Astera Labs 的 Scorpio X 營收＝非 NVIDIA scale-up 交換晶片市場（本專案網通模型：AMD、Trainium、其他 ASIC 顆數 × 每顆 scale-up 交換晶片內容）× Astera 有效份額；PCIe 與訊號調節（Aries 重定時器、Scorpio P、Taurus AEC、Leo CXL）＝AI 加速器顆數 × 每顆內容 × 份額。3Q26 指引季增約 40%、Scorpio 成為最大產品線：成長重心正由重定時器轉到 scale-up 交換晶片。現價約 64 倍共識 FY2026 EPS、40 倍 FY2027。", False),
 ("核心問題：現價隱含的 FY2028 EPS 需要 Scorpio X 在非 NVIDIA scale-up 交換晶片維持多少份額、PCIe 等業務流失多少才成立？非 NVIDIA 加速器顆數（AMD、Trainium、其他 ASIC）與 NVLink Fusion 的競爭是主要不確定性。", False), ("", False),
 ("會計年度：日曆年（12 月底結束）；上游市場層為日曆年，FY(t) 市場＝CY(t)。預測期 FY2027–FY2029（FY2026 為錨點），FY2030–FY2032 為收斂法。口徑：公司 non-GAAP；十年 P/E 規則值用 GAAP 換算 EPS。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("Scorpio X：markets/network 的非 NVIDIA scale-up 交換晶片市場 × 情境顆數係數 × Astera 有效份額（FY2026 錨點校準）× 份額相對值（Assumed）。PCIe 等：AI 加速器顆數 × 情境顆數係數 × 每顆內容 × 有效份額（FY2026 校準）× 相對值（Assumed）。", False),
 ("毛利率：FY2027 起為輸入路徑（Assumed；公司長期目標 70%）。營業費用：彈性 phi（FY2026 錨點 Derived）。稅率：約 12%（公司 non-GAAP 稅率）。股數：+1.5%／年。", False),
 ("校準：FY2026 ＝ 1Q26、2Q26 實績（營收 3.08、3.92 億美元；non-GAAP EPS 0.61、0.80）＋ 3Q26 指引中點（5.50 億美元）＋ 4Q26（季增 15%，Assumed）。產品線拆分公司未揭露：Scorpio X FY2026 為假設，PCIe 等為殘差。不用共識（只在 Recon 對帳）。", False),
 ("估值：目標價（Val 第 0 節）＝0.5 × DCF（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，情境權重 0.3／0.5／0.2）＋0.5 × Comparable（美股高速連接晶片同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）。舊法規則值情境加權保留參考（不是目標價）；市場隱含 FY2028 EPS（Val 第 4 節）；十年 P/E 規則值（只有 2025 年一個正值）× FY2027 GAAP 換算 EPS；循環位置以本模型營收最高年判定。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。客戶名稱公司未具名；10-K 只揭露單一終端客戶逾 70%、前三大約 86%（FY2025）。", False),
 ("", False),
 ("已知限制：(1) Scorpio X FY2026 營收、兩條業務的份額相對值、毛利率路徑皆為 Assumed，是 FY2027–FY2029 EPS 最敏感的輸入（Sens）；(2) 產品線營收未揭露，PCIe 等為殘差；(3) 網通模型沒有 PCIe 重定時器與 PCIe 交換晶片的獨立列，本模型以加速器顆數 × 每顆內容代替；(4) 非 NVIDIA scale-up 交換晶片市場排除 NVIDIA（NVLink）與 TPU（ICI／光學電路交換），NVLink Fusion 若開放第三方交換晶片會改變分母；(5) 十年 P/E 只有 2025 年一個觀測值、含高成長期溢價，規則值明顯偏高；(6) R4 產業部署路徑未建（Recon 預留列）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170


# ================= FY2030–FY2032 收斂法；目標價＝DCF 與 Comparable 加權（與 crdo/ 相同的建置步驟） =================
# 工作單：Andy 2026-10-25 11:18／11:20（三情境延伸到 FY2032、收斂法）、11:28（目標價＝DCF＋Comparable）、11:53（tp_w_dcf 0.5；A 股與其他市場同業分開）、12:22（DCF 下限 0）
import re as _re
from copy import copy as _copy
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
FCX = dict(FC); FCX.update({2030: "H", 2031: "I", 2032: "J"})
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
INR = {IN.cell(i, 1).value: i for i in range(4, IN.max_row + 1) if IN.cell(i, 1).value}
for _k, _i in INR.items():
    if str(_k).startswith(("tp_", "fade_")) or _k in ("cons_tp", "cons_feps"):
        for _j in range(1, 10): IN.cell(_i, _j).fill = KEY

def _move_notes(ws):
    """說明欄（第 8 欄，H）移到第 11 欄（K），讓出 H–J；表頭列補上收斂年度。"""
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 8)
        if c.value is None: continue
        assert not (isinstance(c.value, str) and c.value.startswith("=")), (ws.title, r, c.value)
        k = ws.cell(r, 11); k.value = c.value; k.font = _copy(c.font); k.fill = _copy(c.fill)
        c.value = None; c.font = Font(); c.fill = PatternFill()
        if str(ws.cell(r, 1).value or "") in ("項目", "情境／項目"):
            for j, y in zip((8, 9, 10), FADE):
                h = ws.cell(r, j, f"FY{y}（收斂）"); h.font = BOLD; h.fill = HDR
    ws.column_dimensions["K"].width = ws.column_dimensions["H"].width or 70
    for col in "HIJ": ws.column_dimensions[col].width = 14

def _fade(ws, R, s, eps_f):
    """FY2030–32：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到）；非驅動因子推導。"""
    for y in FADE:
        c, pc, k = FCX[y], FCX[y - 1], y - 2029
        fs = {"g": f"=$G{R['g']}+(fade_g{s}-$G{R['g']})*{k}/3", "rev": f"={pc}{R['rev']}*(1+{c}{R['g']})",
              "nim": f"=$G{R['nim']}+(fade_m{s}-$G{R['nim']})*{k}/3", "ni": f"={c}{R['rev']}*{c}{R['nim']}",
              "sh_": f"={pc}{R['sh_']}*(1+sh_g)", "eps": "=" + eps_f(c, R), "epsg": f"={c}{R['eps']}/{pc}{R['eps']}-1"}
        for key, f in fs.items():
            ws[f"{c}{R[key]}"] = f; ws[f"{c}{R[key]}"].number_format = ws[f"G{R[key]}"].number_format; ws[f"{c}{R[key]}"].fill = FADEFILL

def _shift_rows(wb, sheet, start, n):
    """把 sheet 第 start 列起的內容整體下移 n 列，並同步改寫所有公式與定義名稱中指向這些列的參照（等同插入列）。"""
    ws = wb[sheet]
    cell_ = r"\$?[A-Z]{1,3}\$?\d+"
    pat = _re.compile(r"(?<![A-Za-z0-9_.!$:])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(" + cell_ + r")(?::(" + cell_ + r"))?(?![A-Za-z0-9_(!])")
    one = _re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")
    def fix(f, own):
        parts = f.split('"')
        for i in range(0, len(parts), 2):
            def rep(m):
                pre = m.group(1)
                tgt = (pre is None and own) or (pre is not None and pre.strip("'!") == sheet)
                def sh(x):
                    a, rr = one.fullmatch(x).groups()
                    rr = int(rr)
                    return f"{a}{rr + n if (tgt and rr >= start) else rr}"
                return (pre or "") + sh(m.group(2)) + (":" + sh(m.group(3)) if m.group(3) else "")
            parts[i] = pat.sub(rep, parts[i])
        return '"'.join(parts)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = fix(c.value, w.title == sheet)
    for nm in list(wb.defined_names.keys()):
        d = wb.defined_names[nm]
        if sheet + "!" in d.attr_text: d.attr_text = fix(d.attr_text, False)
    cells = [(c.row, c.column) for row in ws.iter_rows(min_row=start) for c in row if c.value is not None or c.has_style]
    for r, col in sorted(cells, reverse=True):
        src = ws.cell(r, col); dst = ws.cell(r + n, col)
        dst.value = src.value; dst._style = _copy(src._style); dst.number_format = src.number_format
        src.value = None; src._style = _copy(ws.cell(1, 30)._style)

def _tp_section(V, eps_ref, fwd, fwd_note, peers, peer_note, cons_note, method, pe_rng, fmt, cons_feps_note):
    """Val 第 0 節（第 4–48 列）：0.1 DCF、0.2 Comparable、0.3 結論。列標籤依工作單逐字；B／C／D＝情境 A 順利／B 平平／C 不如預期。"""
    T = {}
    def tset(r, key, label, fs, nf, note=""):
        V.cell(r, 1, label); T[key] = r
        for c, f in zip("BCD", fs):
            if f is None: continue
            V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = nf
        if note: V[f"E{r}"] = note
    V.cell(4, 1, "0. 目標價").font = BOLD
    V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
    hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
    V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法（PL 頁灰底）"
    for i, y in enumerate(YRS6):
        tset(8 + i, f"eps{y}", f"  FY{y}", [f"={eps_ref(s, y)}" for s in SC], "0.00")
    V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）")
    tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=IF(tp_roe<=tp_g,0.8,MIN(tp_g/tp_roe,0.8))", None, None], "0.0%", "ROE＝In 頁 tp_roe；tp_roe ≤ tp_g 時取上限 0.8")
    for i, y in enumerate(YRS6):
        tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{8 + i}*(1-$B$15)" for c in "BCD"], "0.00")
    tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], fmt,
         "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027；小於 0 時取 0（股東有限責任）")
    tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], fmt)
    tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], fmt, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
    V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
    hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
    assert 3 <= len(peers) <= 5, peers
    for i, (nm, tk, pe, dt, src) in enumerate(peers, 28):
        V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
    tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{27 + len(peers)})", None, None], "0.0", peer_note)
    tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）",
         ["=" + "+".join(f"tp_w{s}*(" + "+".join(f"{w}*{c}{8 + YRS6.index(y)}" for y, w in fwd) + ")" for s, c in zip(SC, "BCD")), None, None], "0.00", fwd_note)
    tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], fmt, "同業中位數 × 前瞻 EPS")
    V.cell(37, 1, "0.3 結論").font = BOLD
    hdr(V, 38, ["項目", "數值", "", "", "說明"])
    tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], fmt, "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
    tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], fmt)
    tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], fmt, cons_note)
    tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px")
    tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
    tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", cons_feps_note)
    tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
    V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
    tset(46, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D23,"{fmt}")&"～"&TEXT(B23,"{fmt}")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
    tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
         [f"=IF(px/{c}9>MAX({pe_rng}),1,IF(px/{c}9<MIN({pe_rng}),0,PERCENTRANK({pe_rng},px/{c}9)))" for c in "BCD"], "0%",
         "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 6 節採用來源）的分位；高於十年最大值＝100%、低於最小值＝0%；只作檢查，不進目標價")
    V.cell(48, 1, "目標價方法"); V["B48"] = method
    for r_ in (24, 35, 39):
        V[f"B{r_}"].fill = KEY
    return T

def _tp_sens(SE, T):
    """Sens：目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.3／0.7）；放在案例摘要與第 40 列推導區塊之間。"""
    last = max(r for r in range(1, 40) if any(SE.cell(r, j).value is not None for j in range(1, 13)))
    t0 = last + 2
    SE.cell(t0, 1, "目標價參數（只改 DCF 參數與權重；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
    hdr(SE, t0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
    def dcf_f(c, r, g, reinv):
        cf = lambda i: f"Val!{c}{8 + i}*(1-{reinv})"
        pv = "+".join(f"{cf(i)}/(1+{r})^{i + 1}" for i in range(6))
        return f"MAX(0,{pv}+{cf(5)}*(1+{g})/({r}-{g})/(1+{r})^6)"
    cases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
             ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
             ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
    for k, (nm, st, kw) in enumerate(cases):
        i = t0 + 2 + k; r = kw.get("r", "tp_r"); g = kw.get("g", "tp_g"); w = kw.get("w", "tp_w_dcf")
        reinv = f"IF(tp_roe<={g},0.8,MIN({g}/tp_roe,0.8))"
        SE.cell(i, 1, nm); SE.cell(i, 2, st)
        SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_f(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
        SE[f"D{i}"] = f"=Val!B{T['comptp']}"
        SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
        SE[f"F{i}"] = f"=E{i}/E${t0 + 2}-1"
        for c in "CDE": SE[f"{c}{i}"].number_format = "0.00"
        SE[f"F{i}"].number_format = "0.0%"
    assert t0 + 1 + len(cases) < 40, t0
    return t0, t0 + 1 + len(cases)

def _fade_all(eps_f, nim10):
    """PL：三情境推導區塊與頂端摘要延伸到 FY2032；In 頁 fade_m 連到 PL 頁 FY2029 淨利率。nim10＝十年平均淨利率的名稱；None＝未找到，依工作單用 FY2029 × 0.8。"""
    _move_notes(PL)
    for s in SC:
        _fade(PL, PB[s], s, eps_f)
        for key in ("rev", "g", "nim", "ni", "eps", "epsg"):
            if (s, key) not in SUMR: continue
            for y in FADE:
                cc = PL[f"{FCX[y]}{SUMR[(s, key)]}"]; cc.value = f"={FCX[y]}{PB[s][key]}"
                cc.number_format = PL[f"G{SUMR[(s, key)]}"].number_format; cc.fill = FADEFILL
    nimr = {s: f"PL!$G${PB[s]['nim']}" for s in SC}
    IN.cell(INR["fade_mA"], 5, f"={nimr['A']}")
    IN.cell(INR["fade_mB"], 5, f"=({nimr['B']}+{nim10})/2" if nim10 else f"=({nimr['B']}+{nimr['B']}*0.8)/2")
    IN.cell(INR["fade_mC"], 5, f"={nim10}" if nim10 else f"={nimr['C']}*0.8")
    PL["A2"] = ("灰底＝FY2030–FY2032 收斂法，非驅動因子推導（上游市場層只到 2029）：營收年增由 FY2029 線性收斂到 fade_g（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt），"
                "淨利率由 FY2029 線性收斂到 fade_m（In 頁 K 節），FY2032 達到；股數依 sh_g")

def _rule_pe_range(V, ncols=1):
    """第 6 節十年倍數表（2016–2025 列）中採用來源的欄位範圍（PERCENTRANK 用；ncols＞1 時各來源合併成一個分布）。"""
    yr = [r for r in range(1, V.max_row + 1) if V.cell(r, 1).value in [str(y) for y in range(2016, 2026)]]
    yr = [r for r in yr if r > 50][:10]
    assert len(yr) == 10 and yr[-1] - yr[0] == 9, yr
    return f"$B${yr[0]}:${chr(65 + ncols)}${yr[-1]}"

_fade_all(lambda c, R: f"{c}{R['ni']}/{c}{R['sh_']}", None)
_shift_rows(wb, "Val", 4, 47)
TPT = _tp_section(V, lambda s_, y: f"PL!{FCX[y]}{SUMR[(s_, 'eps')]}", [(2027, '2/12'), (2028, '10/12')],
    '約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）',
    [('Credo', 'CRDO', 20.19, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'), ('Marvell', 'MRVL', 40.44, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'),
     ('Broadcom', 'AVGO', 18.52, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'), ('Rambus', 'RMBS', 25.23, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'),
     ('MACOM', 'MTSI', 28.59, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）')],
    '美股高速連接晶片同業（Credo、Marvell、Broadcom、Rambus、MACOM；10-K 列名競爭者為主；不放 A 股）',
    'In 頁 tp_street：stockanalysis 26 家平均 388.36（資料 2026-10-16；190–500）',
    'DCF：股東現金流＝EPS（non-GAAP）×（1 − g／ROE），r＝10%、g＝4%、ROE＝non-GAAP 淨利 ÷ 平均股東權益三期平均（Derived，約 28%；淨現金約 12.5 億美元使 ROE 偏低），FY2027–32 明確期＋永續，FY2030 起為收斂法。Comparable：美股高速連接晶片同業（Credo、Marvell、Broadcom、Rambus、MACOM）前瞻 P/E 中位數 × 12 個月後前瞻 EPS（non-GAAP）；Astera 本身前瞻 P/E 約 46 倍、高於所有同業，同業中位數偏誤方向為低估成長溢價。十年 P/E 只有 2025 年一個觀測值（GAAP），分位檢查參考性低。情境對應：顆數情境 A＝順利、B＝平平、C＝不如預期。',
    _rule_pe_range(V, 2), '#,##0.00',
    "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝In 頁 cons_feps（FY2028 共識為推導值）")
TPS0, TPS_LAST = _tp_sens(SE, TPT)


# ---------------- v0.1.1：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00、14:17）；各頁只在現有最後一列之後追加，不位移既有列 ----------------
_up = next(r for r in range(1, V.max_row + 1) if V.cell(r, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 2; REC_ROW = GRADE_ROW + 1
V.cell(GRADE_ROW, 1, "上檔空間級距")
V.cell(GRADE_ROW, 2, f'=IF(B{_up}>0.5,"大於 50%",IF(B{_up}>0.25,"25%–50%",IF(B{_up}>=0,"0%–25%",IF(B{_up}>=-0.25,"0%～−25%",IF(B{_up}>=-0.5,"−25%～−50%","小於 −50%")))))')
V.cell(GRADE_ROW, 5, f"以「上檔空間：加權目標價」（B{_up}）分六級（Andy 2026-10-25 14:00）")
V.cell(REC_ROW, 1, "投資建議")
V.cell(REC_ROW, 2, f'=IF(B{_up}>0.5,"強力買進",IF(B{_up}>0.25,"買進",IF(B{_up}>=-0.25,"中立",IF(B{_up}>=-0.5,"放空","強力放空"))))')
V.cell(REC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = ['目標價 US$132（DCF 105、Comparable 160 各半），比現價 257.16（2026-10-23）低 48.5%，級距 −25%～−50%；Comparable 單獨為 −37.9%。EPS 與倍數兩個因素都在：本模型 12 個月後前瞻 EPS 6.33 比共識 8.21 低 22.9%；同業（Credo、Marvell、Broadcom、Rambus、MACOM）前瞻 P/E 中位數 25.2 倍，現價隱含 31.3 倍共識 EPS、共識目標價 388.36 隱含 47.3 倍（比同業高 87.5%）；Comparable 未給 Astera 的成長溢價（本身前瞻 P/E 約 46 倍）。', '情境分歧從 FY2028 起：FY2027 EPS 三情境同為 6.90，FY2028 A／B／C 為 8.10／6.13／3.60，DCF 價格帶 52～139；Val 頁讀法：以同業約 25 倍計，現價隱含的 FY2028 EPS 比模型 A 高約三成，需要 Scorpio X 份額上升才達得到。', 'FY2027 EPS 最敏感的是 Scorpio X 份額（高 +18.7%、低 −14.0%）與 PCIe 等份額（高 +15.2%、低 −11.4%）；最大客戶 30% 轉單時 −12.8%（Sens 頁）。下一個檢驗點是 3Q26 財報（2026-11-03）的 Scorpio X 營收與客戶數。']   # 重點觀察（與 README.md 相同文字；數字取自重算後的 Val／PL／Sens 頁）
_rr = RD.max_row + 1
RD.cell(_rr, 1, 'v0.1.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁末新增「上檔空間級距」「投資建議」兩列（Andy 2026-10-25 14:00、14:17 固定規則），README 頁末新增投資建議與重點觀察三點；既有儲存格位置與數值不變。')
RD.cell(_rr + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"').font = Font(bold=True)
for _j, _t in enumerate(OBS, 1): RD.cell(_rr + 1 + _j, 1, f"重點觀察 {_j}：" + _t)
if os.environ.get("RATING_ROWMAP"):   # 除錯用：輸出新增列號（供撰寫 html.json）；CI 不設此變數
    print("ROWMAP", GRADE_ROW, REC_ROW, _rr, _rr + 1, _rr + 2)



# ---------------- 分部原則第 3 條（2026-10-28）：核對列與宣告（追加於工作表最後，不位移既有列，數值不變） ----------------
import os as _os2, sys as _sys2
_sys2.path.insert(0, _os2.path.join(_os2.path.dirname(_os2.path.dirname(_os2.path.abspath(__file__))), "tools", "segment_check"))
from addrow import add_check
add_check(Q, seg_rows=[18, 19], total_row=4, cols="C", rnd=6, note="分部原則第 3 條（2026-10-28 追加）；分部宣告見 segments.json")
RD.cell(RD.max_row + 1, 1, "v0.1.2（2026-10-28）：分部原則重建（數值不變）。Q 頁最後追加核對列「檢查：分部合計 − 合併營收」；官方分部與子列宣告見 segments.json。")
if isinstance(RD.cell(1, 1).value, str): RD.cell(1, 1).value = RD.cell(1, 1).value.replace("v0.1.1", "v0.1.2")

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("ALAB_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "UN_NOTE": UN_NOTE, "CHK": CHK_ROW, "SENSN": SENS_N, "SENS_READ": SENS_READ,
               "VAL": VAL, "IMPL": {str(k): v for k, v in IMPL.items()}, "READ": READ_ROW, "NOPRIOR": NOPRIOR_ROW, "RECON_END": RECON_END, "rule_rows": rule_rows, "SB": SB,
               "TPT": TPT, "TPS": [TPS0, TPS_LAST], "INR": INR},
              open(os.environ["ALAB_ROWMAP"], "w"), ensure_ascii=False, indent=1)
