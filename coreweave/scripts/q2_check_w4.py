# W4（2026-10-08）：Q2 驗證拆解（工作單 W4 第 3 步；只作驗證，不校準 k——已決定事項 14）。
# 模型首期（FY26 下半年）每平均在役 MW 年收入 R_m 對 Q2 實際年化 R_q，拆成：
#   R ＝ B（計費 MW ÷ 在役 MW，爬坡分母）× U（利用率）× A（錨：在役世代 × IF_HoldEcon）× K（定價倍數）＋ O（非算力服務等其他，每在役 MW）
#   Q2：B_q＝季末 Billable ÷ 季末在役 MW [Assumed]、U_q＝模型利用率（Q2 未揭露 [Assumed]）、A_q＝季末世代組合 × IF_HoldEcon、
#       O_q＝R_q × O_m ÷ R_m（服務收入占比同模型 [Assumed]）、K_q＝(R_q − O_q) ÷ (B_q × U_q × A_q)（Q2 隱含 k，調整後）。
#   依序替換（Q2 → 模型）：(i) 爬坡分母 B → (ii) 利用率 U → (iii) 定價倍數 K → (iv) 世代組合 A → (v) 其他 O；各項相加＝總差距（恆等）。
# 數值讀重算後的 Excel「每MW經濟性」頁（C 欄＝首期）。用法：python3 scripts/q2_check_w4.py 檔案.xlsx [輸出 json]
import sys, json
from openpyxl import load_workbook
x = sys.argv[1]
ws = load_workbook(x, data_only=True)['每MW經濟性']
V = {str(ws.cell(r, 1).value): ws.cell(r, 3).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
g = lambda k: float(V[k])
Rm, MWa = g('每 MW 年收入（算力＋服務）'), g('平均在役 MW 合計')
Bm, Um = g('錨｜平均計費 MW') / MWa, g('錨｜利用率')
_KL = next(k for k in ('定價倍數 k（既有占比 × k_既有 ＋（1 − 既有占比）× k_新約）', '定價倍數 k（隨需占比 × k_現貨 ＋（1 − 隨需占比）× k_長約）') if k in V)  # W5 起 k 含既有合約
Am, Km = g('錨｜每 MW 經濟持有成本（IF_HoldEcon，在役世代加權）'), g(_KL)
Om = Rm - Bm * Um * Am * Km
Rq, Bq, Aq = g('最近一季每 MW 年收入（營收 × 4 ÷ 平均在役 MW）'), g('最近一季計費比例（季末 Billable ÷ 季末在役 MW，假設）'), g('最近一季錨（季末在役世代 × IF_HoldEcon）')
Uq, Oq = Um, Rq * Om / Rm
Kq = (Rq - Oq) / (Bq * Uq * Aq)
steps = [('(i) 爬坡分母（計費 MW ÷ 在役 MW）', (Bm - Bq) * Uq * Aq * Kq),
         ('(ii) 利用率', Bm * (Um - Uq) * Aq * Kq),
         ('(iii) 定價倍數 k（模型 k − Q2 隱含 k）', Bm * Um * Aq * (Km - Kq)),
         ('(iv) 世代組合（錨：模型首期平均 vs Q2 季末）', Bm * Um * (Am - Aq) * Km),
         ('(v) 其他（非算力服務等）', Om - Oq)]
tot = Rm - Rq; s = sum(v for _, v in steps)
res = {'model': {'R': Rm, 'B': Bm, 'U': Um, 'A': Am, 'K': Km, 'O': Om}, 'q2': {'R': Rq, 'B': Bq, 'U': Uq, 'A': Aq, 'K': Kq, 'O': Oq, 'kRaw': Rq / Aq},
       'steps': steps, 'total': tot, 'sum': s, 'gapPct': tot / Rq}
print(f'Q2 驗證拆解（US$m／平均在役 MW／年）：模型首期 {Rm:.3f} − Q2 年化 {Rq:.3f} ＝ {tot:+.3f}（{tot / Rq:+.1%}）')
print('| 項目 | 模型 | Q2 | 貢獻 |'); print('|---|---|---|---|')
for (nm, v), (a, b) in zip(steps, [(Bm, Bq), (Um, Uq), (Km, Kq), (Am, Aq), (Om, Oq)]):
    print(f'| {nm} | {a:.3f} | {b:.3f} | {v:+.3f} |')
print(f'| 合計 | {Rm:.3f} | {Rq:.3f} | {s:+.3f} |')
print(f'Q2 隱含 k：未調整 {Rq / Aq:.3f}（年化營收 ÷ 在役 MW ÷ 錨）、調整後 {Kq:.3f}（扣爬坡分母、利用率、服務）；模型首期 k {Km:.3f}')
ok = abs(s - tot) < 0.01
print('拆解檢查：' + ('通過（各項相加＝總差距，誤差 < 0.01）' if ok else f'不通過（{s:.4f} ≠ {tot:.4f}）'))
if len(sys.argv) > 2: json.dump(res, open(sys.argv[2], 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
sys.exit(0 if ok else 1)
