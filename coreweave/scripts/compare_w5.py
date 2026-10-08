# W5（2026-10-08）：v4.6／v4.7 W4／v4.7 W5 三版對照報告（工作單 W5 第 4 步）。
# 讀三版 Excel（LibreOffice 開啟、三情境切換「情境選擇」）的「每MW經濟性」彙總列、加權目標價、評等、融資缺口與 EBITDA 率，
# 新版另讀「公司實況驗證」頁（參數驗證、Q2 逐項對帳；基準情境），加上 attrib_w5.py（變動拆解）與 permw_sens.json（敏感度快照），
# 寫成對照 Excel（數值為讀取值；檢查格為活公式）與 JSON（供 md 報告）。
# 用法：python3 scripts/compare_w5.py --v46 v4.6.xlsx --w4 v4.7_W4.xlsx --w5 v4.7_W5.xlsx --attrib attrib_w5.json --out 報告.xlsx [--json 資料.json]
import sys, os, json, argparse
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
from uno_q import Workbook
from openpyxl import Workbook as XW
from openpyxl.styles import Font, PatternFill

SUM = ['每 MW 年收入（算力＋服務）', '電費', 'IT 維護', '人員、軟體、水與耗材', '財產稅與保險', '公司管銷與其他', '租金', '現金成本合計（含租金）', 'EBITDA',
       'D&A（模型車隊折舊）', '利息', '稅前', '每 MW 資本支出（新增 MW 的建置成本）']
KROWS = ['定價倍數 k（既有占比 × k_既有 ＋（1 − 既有占比）× k_新約）', '定價倍數 k（隨需占比 × k_現貨 ＋（1 − 隨需占比）× k_長約）',
         '既有合約占比（平均既有 MW ÷ 平均在役 MW × 開關）', 'k_既有（Q2 實現單價 ÷ Q2 錨；合約期內固定）', 'k_新約（隨需占比 × k_現貨 ＋（1 − 隨需占比）× k_長約 ×（1＋新約價格調整））',
         '錨｜每 MW 經濟持有成本（IF_HoldEcon，在役世代加權）', 'Tokenomics 錨 × k：每 MW 年收入（100% 計費時數）']
SCN = {1: '保守 4.2 GW', 2: '基準 5.6 GW', 3: '積極 8 GW'}
CODE = {1: '買進', 0: '中立', -1: '賣出'}
P = ['FY26（2H）', 'FY27', 'FY28', 'FY29', 'FY30']


def grab(x):
    out = {}
    with Workbook(x) as wb:
        sel = wb.cell('輸入與假設', '情境選擇', prefix=True); s0 = wb.get(sel)
        def row(sh, nm):
            try: return [wb.get(wb.cell(sh, nm, col=c)) for c in 'CDEFG']
            except KeyError: return None
        for s in (1, 2, 3):
            wb.set(sel, s); o = {nm: row('每MW經濟性', nm) for nm in SUM + KROWS}
            o['tp'] = wb.get(wb.cell('評價_DCF與目標價', '加權目標價'))
            o['rating'] = CODE[round(wb.get(wb.cell('評價_DCF與目標價', '目標價區間｜點位評等代碼（1＝買進、0＝中立、−1＝賣出）')))]
            o['gap'] = max(0.0, -wb.get(wb.cell('各期收支', '融資前累積現金', col='G')))
            o['ebm'] = row('運營_產能與收入', 'EBITDA 率（損益與資金共用）')
            out[s] = o
        wb.set(sel, s0)
    return out


def cvtab(x):  # 「公司實況驗證」頁讀值（重算後存檔的快取值＝基準情境）：A–G 數值與文字、I–M 說明
    from openpyxl import load_workbook
    ws = load_workbook(x, data_only=True)['公司實況驗證']
    return [[ws.cell(r, c).value for c in (1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13)] for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value]


def main():
    ap = argparse.ArgumentParser()
    for k in ('v46', 'w4', 'w5', 'attrib', 'out'): ap.add_argument('--' + k, required=True)
    ap.add_argument('--json'); a = ap.parse_args()
    A, B, C = grab(os.path.abspath(a.v46)), grab(os.path.abspath(a.w4)), grab(os.path.abspath(a.w5)); C[2]['cv'] = cvtab(os.path.abspath(a.w5))
    AT = json.load(open(a.attrib, encoding='utf-8'))
    SN = json.load(open(os.path.join(ROOT, 'permw_sens.json'), encoding='utf-8'))
    data = {'v46': A, 'w4': B, 'w5': C, 'attrib': AT, 'sens': SN}
    H, T, S_ = Font(bold=True, color='FFFFFF'), Font(bold=True, size=13), Font(size=9, color='555555')
    FH = PatternFill('solid', fgColor='1F3864')
    wb = XW(); ws = wb.active; ws.title = '摘要'
    def hdr(ws, r, xs):
        for j, x in enumerate(xs):
            c = ws.cell(row=r, column=1 + j, value=x); c.font = H; c.fill = FH
    ws['A1'] = 'CoreWeave v4.6 → v4.7 W4 → v4.7 W5（公司實況驗證）三版對照'; ws['A1'].font = T
    ws['A2'] = '公司調整只採用有證據的機制（既有合約以 Q2 實現單價）；營運成本與每 MW 資本支出差距 > 10% 但找不到機制：Tokenomics 為基準、公司實際為敏感度'; ws['A2'].font = S_
    r = 4; hdr(ws, r, ['情境', 'v4.6 目標價', 'v4.7 W4', 'v4.7 W5', 'W5 − W4', 'v4.6 評等', 'W4 評等', 'W5 評等', 'v4.6 融資缺口', 'W4 融資缺口', 'W5 融資缺口']); r += 1
    for s in (1, 2, 3):
        for j, v in enumerate([SCN[s], A[s]['tp'], B[s]['tp'], C[s]['tp'], f'=D{r}-C{r}', A[s]['rating'], B[s]['rating'], C[s]['rating'], A[s]['gap'], B[s]['gap'], C[s]['gap']]):
            ws.cell(row=r, column=1 + j, value=v).number_format = '0.00'
        r += 1
    # 每 MW 三版（基準）
    w = wb.create_sheet('每MW_三版'); w['A1'] = '基準情境每 MW（US$m／平均在役 MW／年）：v4.6／v4.7 W4／v4.7 W5'; w['A1'].font = T
    r = 3; hdr(w, r, ['項目', '版本'] + P); r += 1
    for nm in SUM:
        for tag, X in (('v4.6', A), ('v4.7 W4', B), ('v4.7 W5', C)):
            w.cell(row=r, column=1, value=nm); w.cell(row=r, column=2, value=tag)
            for i, v in enumerate(X[2][nm] or []): w.cell(row=r, column=3 + i, value=v).number_format = '0.000'
            r += 1
    for tag, X in (('v4.6', A), ('v4.7 W4', B), ('v4.7 W5', C)):
        w.cell(row=r, column=1, value='EBITDA 率（模型採用）'); w.cell(row=r, column=2, value=tag)
        for i, v in enumerate(X[2]['ebm']): w.cell(row=r, column=3 + i, value=v).number_format = '0.0%'
        r += 1
    r += 1; w.cell(row=r, column=1, value='定價倍數 k 與既有合約（基準；v4.7 W5）').font = Font(bold=True); r += 1
    for nm in KROWS:
        v = C[2][nm]
        if v is None: continue
        w.cell(row=r, column=1, value=nm)
        for i, x in enumerate(v): w.cell(row=r, column=3 + i, value=x).number_format = '0.000'
        r += 1
    # 公司實況驗證（基準情境讀值）
    w = wb.create_sheet('公司實況驗證'); w['A1'] = '公司實況驗證（v4.7 W5 Excel「公司實況驗證」頁的讀值；基準情境）'; w['A1'].font = T
    r = 3; hdr(w, r, ['參數', '單位', 'Tokenomics 值', 'CRWV 實際', '差距', '採用值', '規則', 'Tokenomics 名稱', 'CRWV 實際與標記', '差距原因（機制與證據）', '公司調整', '證據']); r += 1
    for x in C[2]['cv']:
        for j, v in enumerate(x): w.cell(row=r, column=1 + j, value=v)
        r += 1
    # 敏感度
    w = wb.create_sheet('敏感度'); w['A1'] = '敏感度（建置時快照 permw_sens.json；加權目標價 US$／融資缺口 US$bn）'; w['A1'].font = T
    r = 3; hdr(w, r, ['設定', '保守 目標價', '保守 缺口', '基準 目標價', '基準 缺口', '積極 目標價', '積極 缺口']); r += 1
    for k, nm in SN['cases']:
        w.cell(row=r, column=1, value=nm)
        for j, sk in enumerate(('low', 'base', 'high')):
            v = SN['scenarios'][sk].get(k)
            w.cell(row=r, column=2 + 2 * j, value='不適用' if v is None else v[0]); w.cell(row=r, column=3 + 2 * j, value='不適用' if v is None else v[1])
        r += 1
    # 變動拆解
    w = wb.create_sheet('變動拆解'); w['A1'] = '目標價變動拆解（v4.6 → v4.7 W5；0 截斷口徑，US$／股；scripts/attrib_w5.py）'; w['A1'].font = T
    r = 3; hdr(w, r, ['(d) 逐項'] + list(AT['scenarios'])); r += 1
    top = r
    for k, nm in AT['steps']:
        w.cell(row=r, column=1, value=nm)
        for j, sc in enumerate(AT['scenarios']): w.cell(row=r, column=2 + j, value=AT['scenarios'][sc]['seq'][k]['blend']).number_format = '+0.00;-0.00'
        r += 1
    w.cell(row=r, column=1, value='合計（＝總變動）')
    for j in range(3): w.cell(row=r, column=2 + j, value=f'=SUM({chr(66 + j)}{top}:{chr(66 + j)}{r - 1})').number_format = '+0.00;-0.00'
    r += 1; w.cell(row=r, column=1, value='總變動（v4.7 W5 − v4.6，讀值）')
    for j, sc in enumerate(AT['scenarios']): w.cell(row=r, column=2 + j, value=AT['scenarios'][sc]['total']).number_format = '+0.00;-0.00'
    r += 1; w.cell(row=r, column=1, value='檢查（合計 − 總變動，應為 0）')
    for j in range(3): w.cell(row=r, column=2 + j, value=f'=IF(ABS({chr(66 + j)}{r - 2}-{chr(66 + j)}{r - 1})<0.01,"通過","不通過")')
    wb.save(a.out); print('saved', a.out)
    if a.json: json.dump(data, open(a.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)


if __name__ == '__main__':
    sys.exit(main())
