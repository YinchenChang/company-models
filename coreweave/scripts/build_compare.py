# W3（2026-10-08）：v4.5 → v4.6 前後對照 Excel（工作單 W3 第 4 步；Andy 以 Excel 閱讀）。
# 輸入：scripts/compare_gather.py 的取數 JSON（v4.5 成品、舊方法副本、v4.6 成品 × 三情境）、scripts/attrib_permw.py 的拆解 JSON、
#       company.json（Q2 10-Q 數字、世代組合）、Tokenomics 快照。
# 數值為建置時取自兩版成品的值（藍字＝取自成品，「來源」欄寫工作表｜列名稱）；差異、差異 %、Q2 每 MW 年化、設施口徑換算、拆解合計為活公式。
# 用法：python3 scripts/build_compare.py 取數.json 拆解.json 輸出.xlsx [摘要.json]
import sys, os, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
D = json.load(open(sys.argv[1], encoding='utf-8')); AT = json.load(open(sys.argv[2], encoding='utf-8')); OUT = sys.argv[3]
SUMJ = sys.argv[4] if len(sys.argv) > 4 else None
CO = json.load(open(os.path.join(ROOT, 'company.json'), encoding='utf-8'))
TK = json.load(open(os.path.join(ROOT, CO['tokenomics']['snapshotFile']), encoding='utf-8'))
import subprocess
PER = json.loads(subprocess.run(['python3', os.path.join(ROOT, 'calendar_q.py'), ROOT], capture_output=True, text=True).stdout)['periods']
SCN = [('1', '保守 4.2 GW'), ('2', '基準 5.6 GW'), ('3', '積極 8 GW')]
SCK = {'1': '保守', '2': '基準', '3': '積極'}
PM = '每MW經濟性|'
g = lambda tag, s, k: D[tag][s][k]

TITLE = Font(name='Arial', size=13, bold=True); BOLD = Font(name='Arial', size=10, bold=True); BLACK = Font(name='Arial', size=10)
BLUE = Font(name='Arial', size=10, color='0000FF'); SMALL = Font(name='Arial', size=9, color='666666'); HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FH = PatternFill('solid', fgColor='1F3864'); FS = PatternFill('solid', fgColor='D9E2F3'); FY_ = PatternFill('solid', fgColor='FFF2CC')
BOX = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
NUM, NUM1, PCT, USD, SH = '#,##0.00;(#,##0.00);-', '#,##0.0;(#,##0.0);-', '0.0%;(0.0%);-', '$#,##0.00;($#,##0.00);-', '0.000'
wb = Workbook()


def sheet(name, widths):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths): ws.column_dimensions[CL(i + 1)].width = w
    return ws


def hdr(ws, r, cols):
    for j, h in enumerate(cols):
        c = ws.cell(row=r, column=1 + j, value=h); c.font = HEAD; c.fill = FH; c.alignment = Alignment(wrap_text=True, vertical='center')
    return r + 1


def sec(ws, r, t, n=10):
    for j in range(n): ws.cell(row=r, column=1 + j).fill = FS
    ws.cell(row=r, column=1, value=t).font = BOLD
    return r + 1


def put(ws, r, c, v, fmt=None, font=BLACK, wrap=False):
    x = ws.cell(row=r, column=c, value=v); x.font = font; x.border = BOX
    if fmt: x.number_format = fmt
    if wrap: x.alignment = Alignment(wrap_text=True, vertical='top')
    return x


# ---------------- 數字 ----------------
B = '2'
pm = lambda tag, s, k, i: g(tag, s, PM + k)[i]
tgt = lambda tag, s: g(tag, s, '評價_DCF與目標價|加權目標價')[0]
gap = lambda tag, s: max(0.0, -g(tag, s, '各期收支|融資前累積現金')[4])
nsh = lambda tag, s, i: g(tag, s, '各期收支|累計新股')[i]
LQ = CO['latestQuarter']; QMW = CO['fleet']['openMix']['activeMW'] - CO['callFacts']['activeAddQ2'] / 2
att = AT['scenarios']

# ======================= 摘要 =======================
ws = wb.active; ws.title = '摘要'
for i, w in enumerate([34, 10, 8, 11, 11, 11, 10, 70, 44]): ws.column_dimensions[CL(i + 1)].width = w
ws['A1'] = 'CoreWeave v4.5 → v4.6 前後對照（每 MW 改接 Tokenomics v5.26）'; ws['A1'].font = TITLE
b5, b6 = tgt('v45', B), tgt('v46', B)
c27a = pm('v45pm', B, '現金成本合計（含租金）', 1); c27b = pm('v46', B, '現金成本合計（含租金）', 1)
c30a = pm('v45pm', B, '現金成本合計（含租金）', 4); c30b = pm('v46', B, '現金成本合計（含租金）', 4)
CONC = (f"結論：每 MW 年收入不變（收入採備案，仍為 v4.5 輸入）；每 MW 現金成本（含租金）FY27 {c27a:.2f} → {c27b:.2f}、FY30 {c30a:.2f} → {c30b:.2f}（US$m／MW／年），"
        f"每 MW 資本支出 34 → 37.5–37.6；主因是營運成本改由下而上（Tokenomics IT 維護約 1.0、人員軟體 0.33、稅險 0.17 加電費 0.67）與 IF_CapexIT 建置成本，"
        f"基準目標價 ${b5:.2f} → ${b6:.2f}（{b6 - b5:+.2f}），三情境皆賣出。")
ws['A2'] = CONC; ws['A2'].font = BOLD; ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:I2'); ws.row_dimensions[2].height = 48
ws['A3'] = '藍字＝建置時取自兩版成品（v4.5：dist v4.5 與舊方法副本，逐格相同；v4.6：dist v4.6），黑字＝公式。基準情境；每 MW＝每平均在役 MW（IT 關鍵電力）、年化。'; ws['A3'].font = SMALL
r = hdr(ws, 5, ['項目', '單位', '年度', 'v4.5', 'v4.6', '差異', '差異 %', '主要原因（一句）', '來源（成品工作表｜列名稱）'])
SUM = []  # 給 md 與 PR 用


def srow(name, unit, fy, a, b, why, src, fmt=NUM):
    global r
    put(ws, r, 1, name, font=BOLD if not name.startswith('　') else BLACK); put(ws, r, 2, unit, font=SMALL); put(ws, r, 3, fy)
    put(ws, r, 4, a, fmt, BLUE if a is not None else SMALL); put(ws, r, 5, b, fmt, BLUE)
    if a is not None:
        put(ws, r, 6, f'=E{r}-D{r}', fmt); put(ws, r, 7, f'=IF(ABS(D{r})<1E-9,"—",E{r}/D{r}-1)', PCT)
    else:
        put(ws, r, 6, '—（v4.5 未拆項）', font=SMALL); put(ws, r, 7, None)
    put(ws, r, 8, why, wrap=True); put(ws, r, 9, src, font=SMALL, wrap=True)
    SUM.append({'name': name, 'unit': unit, 'fy': fy, 'v45': a, 'v46': b, 'why': why})
    r += 1


ITEM = [('電費', 'IF_PowerCost（含平均用電比與 PUE）'), ('IT 維護', 'IF_MaintIT 世代加權（GB300 1.12、VR200 1.13；W2 暫代 0.15）'),
        ('人員、軟體、水與耗材', 'IF_StaffSW 0.325（全額）'), ('財產稅與保險', 'IF_TaxIns × IT 資本占比（只算 CRWV 擁有的 IT）'),
        ('公司管銷與其他', '管銷率 5.51%（Q2 10-Q，扣 SBC）× 營收'), ('租金', '未改（10-Q 租約到期表＋表外租金路徑）')]
for fi, fy in ((1, PER[1]), (4, PER[4])):
    srow('每 MW 年收入（算力＋服務）', 'US$m/MW', fy, pm('v45pm', B, '每 MW 年收入（算力＋服務）', fi), pm('v46', B, '每 MW 年收入（算力＋服務）', fi),
         '未改：收入採備案（每 MW 年收入仍為 v4.5 輸入 11.2–10.5 × 利用率；GPU 小時長約價不足兩個獨立來源）', '每MW經濟性｜每 MW 年收入（算力＋服務）')
    srow('每 MW 現金成本（含租金）', 'US$m/MW', fy, pm('v45pm', B, '現金成本合計（含租金）', fi), pm('v46', B, '現金成本合計（含租金）', fi),
         ('由下而上（Tokenomics 電費、IT 維護、人員軟體、稅險＋管銷＋租金）取代 v4.5 的 EBITDA 率線性路徑（59% → 65%）' if fi == 1 else
          'v4.5 穩態 65% 讓隱含成本隨收入下降；由下而上每 MW 成本隨世代略升（IT 維護、稅險），差距擴大'),
         'v4.5＝收入 − EBITDA（隱含）；v4.6＝每MW經濟性｜現金成本合計（含租金）')
    srow('　非租金現金成本', 'US$m/MW', fy, pm('v45pm', B, '現金成本合計（含租金）', fi) - pm('v45pm', B, '租金', fi),
         pm('v46', B, '現金成本合計（含租金）', fi) - pm('v46', B, '租金', fi), 'v4.5 隱含（未拆項）；v4.6＝下列 5 項合計', '＝現金成本合計 − 租金')
    for k, why in ITEM:
        srow('　' + k, 'US$m/MW', fy, pm('v45pm', B, k, fi) if k == '租金' else None, pm('v46', B, k, fi), why, '每MW經濟性｜' + k)
    srow('每 MW EBITDA', 'US$m/MW', fy, pm('v45pm', B, 'EBITDA', fi), pm('v46', B, 'EBITDA', fi), '＝收入 − 現金成本（含租金）；EBITDA 率 '
         + f"{pm('v45pm', B, 'EBITDA', fi) / pm('v45pm', B, '每 MW 年收入（算力＋服務）', fi):.1%} → {pm('v46', B, 'EBITDA', fi) / pm('v46', B, '每 MW 年收入（算力＋服務）', fi):.1%}", '每MW經濟性｜EBITDA')
    srow('每 MW D&A', 'US$m/MW', fy, pm('v45pm', B, 'D&A（模型車隊折舊）', fi), pm('v46', B, 'D&A（模型車隊折舊）', fi),
         '建置成本提高 → 成長型 CapEx 與 PP&E 增加；折舊年限 IF_DeprLifeIT 加權仍 6 年（不變）', '每MW經濟性｜D&A（模型車隊折舊）')
    srow('每 MW 資本支出（新增 MW）', 'US$m/MW', fy, pm('v45pm', B, '每 MW 資本支出（新增 MW 的建置成本）', fi), pm('v46', B, '每 MW 資本支出（新增 MW 的建置成本）', fi),
         'Σ 新增世代占比 × IF_CapexIT（GB300 37.4、VR200 37.6；不含廠房）取代以 FY26 CapEx 指引校準的 34', '輸入與假設｜每 MW 建置成本')
for s, nm in SCN:
    a = att[SCK[s]]['seq']
    srow(f'加權目標價｜{nm}', 'US$', '—', tgt('v45', s), tgt('v46', s),
         f"(d) 方法變更：① 每 MW 資本支出 {a['capex']['blend']:+.2f}、③ 營運成本由下而上 {a['cost']['blend']:+.2f}（② 折舊年限、④ 收入、⑤ MW 口徑皆 0）", '評價_DCF與目標價｜加權目標價', USD)
for s, nm in SCN:
    srow(f'融資缺口（FY30 累積）｜{nm}', 'US$bn', PER[4], gap('v45', s), gap('v46', s), '模型期 CapEx 增加（建置成本 +10%）且 EBITDA 減少（由下而上成本）',
         '各期收支｜融資前累積現金（FY30，取負值）')
for fi in (1, 4):
    srow('累計新股（基準）', 'bn 股', PER[fi], nsh('v45', B, fi), nsh('v46', B, fi), '融資缺口擴大 → 股權募資（每年上限＝現市值 20%）提早觸頂，其餘以高息債補足',
         '各期收支｜累計新股', SH)

# ---- Q2 2026 實際對帳 ----
r += 1
r = sec(ws, r, f"Q2 2026 實際對帳：模型 {PER[0]}（下半年，基準）vs 10-Q（{LQ['periodEnd']} 季度；每平均在役 MW 年化）", 9)
ws.cell(row=r, column=1, value=f"10-Q 輸入（US$bn，三個月；{LQ['filed']} 申報）").font = BOLD; r += 1
QIN = {}
for k, lab, v in [('rev', '營收', LQ['revenue']), ('cor', '營收成本', LQ['costRev']), ('ti', '技術與基礎設施', LQ['techInfra']), ('sm', '銷售行銷', LQ['sm']),
                  ('smS', '　其中 SBC', LQ['smSbc']), ('ga', '一般管理', LQ['ga']), ('gaS', '　其中 SBC', LQ['gaSbc']), ('da', 'D&A 合計（未依列別揭露；全數視為在營收成本與技術基礎設施）', LQ['da']),
                  ('ctS', 'SBC（營收成本＋技術與基礎設施）', LQ['sbcCostTi']), ('op', '營業租賃成本', LQ['opLeaseCost']), ('var', '變動租賃成本', LQ['varLeaseCost']),
                  ('adj', 'Adj. EBITDA（公司揭露）', LQ['adjEbitda']), ('mw', '平均在役 MW（＝6/30 主動電力 1,500 − Q2 新增 500 ÷ 2）', QMW)]:
    put(ws, r, 1, lab); put(ws, r, 4, v, '#,##0' if k == 'mw' else '0.000', BLUE); QIN[k] = f'$D${r}'; r += 1
r = hdr(ws, r, ['項目', '單位', '', 'Q2 實際（年化）', f'模型 {PER[0]}', '差距（模型 − 實際）', '差距 %', '差距原因', 'Q2 算式'])
Q = QIN; ann = f"*4/{Q['mw']}*1000"
QROWS = [
    ('每 MW 年收入', f"={Q['rev']}{ann}", pm('v46', B, '每 MW 年收入（算力＋服務）', 0),
     '屬 v4.5 既有假設、非本次改動：模型收入以「平均在役 MW × 每 MW 年收入 11.2 × 利用率 95%」＋服務計；Q2 分母含 Q2 新增、尚在爬坡未計費的 MW', '營收 × 4 ÷ 平均在役 MW'),
    ('每 MW 現金成本（含租金、含管銷）', f"=({Q['rev']}-{Q['adj']}){ann}", pm('v46', B, '現金成本合計（含租金）', 0),
     '下列三項合計：營運成本高、租金低、管銷略高', '（營收 − Adj. EBITDA）× 4 ÷ MW'),
    ('　營運成本（租金前、不含管銷）', f"=({Q['cor']}+{Q['ti']}-{Q['da']}-{Q['ctS']}-{Q['op']}-{Q['var']}){ann}",
     pm('v46', B, '電費', 0) + pm('v46', B, 'IT 維護', 0) + pm('v46', B, '人員、軟體、水與耗材', 0) + pm('v46', B, '財產稅與保險', 0),
     'Tokenomics 是滿載穩態口徑（IT 維護約 IT 資本 3%／年、人員軟體 0.33、稅險 0.16）；CRWV 機隊多在原廠保固期、Q2 分母含爬坡 MW、部分電費可能含在變動租賃（房東轉嫁）、D&A 列別未揭露（假設全在此兩列）',
     '（營收成本＋技術與基礎設施 − D&A − SBC − 營業與變動租賃）× 4 ÷ MW'),
    ('　租金（營業＋變動）', f"=({Q['op']}+{Q['var']}){ann}", pm('v46', B, '租金', 0),
     '屬 v4.5 既有租金路徑、非本次改動：模型用 10-Q 租約到期表＋表外租金；Q2 含變動租賃 0.15（可能含電費轉嫁）', '（營業＋變動租賃成本）× 4 ÷ MW'),
    ('　管銷（扣 SBC）', f"=({Q['sm']}-{Q['smS']}+{Q['ga']}-{Q['gaS']}){ann}", pm('v46', B, '公司管銷與其他', 0),
     '同一管銷率 5.51%，模型每 MW 收入較高故金額較高', '（銷售行銷 − SBC ＋ 一般管理 − SBC）× 4 ÷ MW'),
    ('每 MW EBITDA', f"={Q['adj']}{ann}", pm('v46', B, 'EBITDA', 0), '收入高 1.8、成本高 0.8 → EBITDA 高 1.0', 'Adj. EBITDA × 4 ÷ MW'),
    ('EBITDA 率', f"={Q['adj']}/{Q['rev']}", pm('v46', B, 'EBITDA', 0) / pm('v46', B, '每 MW 年收入（算力＋服務）', 0),
     '兩者接近（−0.9pt）是組成互相抵銷的結果：收入與營運成本都高於 Q2、租金低於 Q2；不代表各項已校準', 'Adj. EBITDA ÷ 營收'),
]
QS = []
for nm, fq, mv, why, how in QROWS:
    isP = nm == 'EBITDA 率'
    put(ws, r, 1, nm, font=BOLD if not nm.startswith('　') else BLACK); put(ws, r, 2, '%' if isP else 'US$m/MW', font=SMALL)
    put(ws, r, 4, fq, PCT if isP else NUM); put(ws, r, 5, mv, PCT if isP else NUM, BLUE)
    put(ws, r, 6, f'=E{r}-D{r}', PCT if isP else NUM); put(ws, r, 7, None if isP else f'=E{r}/D{r}-1', None if isP else PCT)
    put(ws, r, 8, why, wrap=True); put(ws, r, 9, how, font=SMALL, wrap=True)
    QS.append({'name': nm, 'model': mv, 'why': why, 'how': how}); r += 1
# Q2 數值（給 md）：以 Python 同式計算
qv = {'每 MW 年收入': LQ['revenue'] * 4 / QMW * 1e3, '每 MW 現金成本（含租金、含管銷）': (LQ['revenue'] - LQ['adjEbitda']) * 4 / QMW * 1e3,
      '　營運成本（租金前、不含管銷）': (LQ['costRev'] + LQ['techInfra'] - LQ['da'] - LQ['sbcCostTi'] - LQ['opLeaseCost'] - LQ['varLeaseCost']) * 4 / QMW * 1e3,
      '　租金（營業＋變動）': (LQ['opLeaseCost'] + LQ['varLeaseCost']) * 4 / QMW * 1e3, '　管銷（扣 SBC）': (LQ['sm'] - LQ['smSbc'] + LQ['ga'] - LQ['gaSbc']) * 4 / QMW * 1e3,
      '每 MW EBITDA': LQ['adjEbitda'] * 4 / QMW * 1e3, 'EBITDA 率': LQ['adjEbitda'] / LQ['revenue']}
for q in QS: q['q2'] = qv[q['name']]
ws.cell(row=r, column=1, value='讀法：模型 FY26 EBITDA 率與 Q2 實際相近，但每 MW 收入、營運成本、租金三項各自差距大，方向相反而抵銷；本次（W3）改動的只有營運成本一項，收入與租金差距屬 v4.5 既有假設。').font = SMALL
ws.freeze_panes = 'B6'

# ======================= 每MW_前後 =======================
ws = sheet('每MW_前後', [40, 10] + [9] * 15 + [50])
ws['A1'] = f'每 MW 經濟性 v4.5 vs v4.6（{PER[0]}–{PER[4]} × 三情境；US$m／MW／年；W2 第 7 步彙總表各列）'; ws['A1'].font = TITLE
ws['A2'] = ('v4.5 欄取自舊方法副本的「每MW經濟性」頁（與 v4.5 成品逐格相同，見 compare_gather 檢查）；電費、IT 維護、人員軟體、稅險、管銷在 v4.5 不入損益（EBITDA 率為輸入），'
            '列為「—」。設施口徑區＝IT 口徑值 × IT MW ÷ 設施 MW（公式）。'); ws['A2'].font = SMALL
ROWS = ["平均在役 MW（IT 關鍵電力）", "平均在役 MW（設施＝IT × IF_FacilityGW）", "每 MW 年收入（算力＋服務）", "　其中：算力收入", "電費", "IT 維護", "人員、軟體、水與耗材", "財產稅與保險",
        "公司管銷與其他", "租金", "現金成本合計（含租金）", "EBITDA", "D&A（模型車隊折舊）", "利息", "稅前", "每 MW 資本支出（新增 MW 的建置成本）", "每設施 MW 年收入", "每設施 MW EBITDA"]
NOV45 = {"電費", "IT 維護", "人員、軟體、水與耗材", "財產稅與保險", "公司管銷與其他"}
r = 4
for s, nm in SCN:
    r = sec(ws, r, f'{nm}', 18)
    r = hdr(ws, r, ['項目（IT 口徑）', '單位'] + [f'v4.5 {p}' for p in PER] + [f'v4.6 {p}' for p in PER] + [f'差異 {p}' for p in PER] + ['說明'])
    top = r
    for k in ROWS:
        put(ws, r, 1, k, font=BOLD if k in ('每 MW 年收入（算力＋服務）', '現金成本合計（含租金）', 'EBITDA', '稅前') else BLACK)
        put(ws, r, 2, 'MW' if k.startswith('平均在役') else 'US$m/MW', font=SMALL)
        fm = '#,##0' if k.startswith('平均在役') else NUM
        for i in range(5):
            if k in NOV45: put(ws, r, 3 + i, '—', font=SMALL)
            else: put(ws, r, 3 + i, pm('v45pm', s, k, i), fm, BLUE)
            put(ws, r, 8 + i, pm('v46', s, k, i), fm, BLUE)
            put(ws, r, 13 + i, '—' if k in NOV45 else f'={CL(8 + i)}{r}-{CL(3 + i)}{r}', fm, SMALL if k in NOV45 else BLACK)
        if k in NOV45: put(ws, r, 18, 'v4.5 未拆項（EBITDA 率為輸入；現金成本＝收入 − EBITDA）', font=SMALL)
        if k == '現金成本合計（含租金）': put(ws, r, 18, 'v4.5＝收入 − EBITDA（隱含）', font=SMALL)
        r += 1
    mi, mf = top, top + 1
    r = hdr(ws, r, ['設施口徑（每設施 MW；＝IT 值 × IT MW ÷ 設施 MW）', '單位'] + [f'v4.5 {p}' for p in PER] + [f'v4.6 {p}' for p in PER] + [f'差異 {p}' for p in PER] + ['說明'])
    for k in ("每 MW 年收入（算力＋服務）", "現金成本合計（含租金）", "EBITDA", "D&A（模型車隊折舊）", "利息", "稅前", "每 MW 資本支出（新增 MW 的建置成本）"):
        src = top + ROWS.index(k)
        put(ws, r, 1, '每設施 MW｜' + k); put(ws, r, 2, 'US$m/MW', font=SMALL)
        for i in range(5):
            for c0 in (3, 8):
                put(ws, r, c0 + i, f'={CL(c0 + i)}{src}*{CL(c0 + i)}{mi}/{CL(c0 + i)}{mf}', NUM)
            put(ws, r, 13 + i, f'={CL(8 + i)}{r}-{CL(3 + i)}{r}', NUM)
        r += 1
    r += 1
ws.freeze_panes = 'C4'

# ======================= 參數對照 =======================
ws = sheet('參數對照', [34, 12, 22, 34, 26, 40, 14, 60])
ws['A1'] = '參數對照：v4.6 改動的每一個參數（舊值｜新值｜來源｜標記）'; ws['A1'].font = TITLE
ws['A2'] = 'Tokenomics 值取自快照 ' + CO['tokenomics']['snapshotFile'] + f"（{TK['source']['file']}，commit {TK['source']['commit'][:7]}，擷取 {TK['source']['extractedAt']}）；主值＝基準成本情境，低／高只作敏感度。每 GW 值＝每 IT MW 百萬美元。"; ws['A2'].font = SMALL
r = hdr(ws, 4, ['參數', '單位', '舊值（v4.5）', '新值（v4.6，基準）', '低／高（敏感度區間）', 'Tokenomics 名稱或公司來源', '標記', '說明'])
GEN = CO['fleet']['generations']; GC = {x['name']: x['code'] for x in TK['source']['generations']}
tkv = lambda n, gg, cs='基準': TK['items'][n]['values'][gg][cs]
gl = lambda n, f='{:.3f}': '；'.join(f"{GC[gg]} {f.format(tkv(n, gg))}" for gg in GEN if gg in TK['items'][n]['values'])
gr = lambda n, f='{:.3f}': '；'.join(f"{GC[gg]} {f.format(tkv(n, gg, '低成本'))}–{f.format(tkv(n, gg, '高成本'))}" for gg in GEN)
v46 = lambda k: '、'.join(f'{x:.2f}' for x in g('v46', B, k))
PAR = [
    ('每 MW 建置成本（新增 MW）', 'US$m/MW', '34（各期）', '、'.join(f"{x:.2f}" for x in g('v46', B, '輸入與假設|每 MW 建置成本')),
     gr('IF_CapexIT', '{:.1f}'), 'Σ 新增世代占比 × IF_CapexIT（不含廠房 IF_CapexFacility；廠房在租金）', '[Derived]', f"IF_CapexIT 基準：{gl('IF_CapexIT', '{:.1f}')}。舊值以 FY26 CapEx 指引回推（W2 起只作對照）"),
    ('GPU 經濟壽命（折舊年限）', '年', '6', f"{g('v46', B, '輸入與假設|GPU 經濟壽命（年）')[0]:g}", '高成本 4', 'IF_DeprLifeIT 依期初在役世代加權、取整數', '[Derived]', f"基準各世代 6 年；與舊值相同，數字不變（拆解 ② ＝ 0）"),
    ('EBITDA 率路徑', '%', '59% → 65%（線性）', '、'.join(f"{x:.1%}" for x in g('v46', B, '輸入與假設|EBITDA 率')), '—', 'methodology.perMw.cost＝bottomUp（每MW經濟性頁由下而上推導）', '[Derived]', '起始 59%（Q2 實際）與穩態 65% 改為對照；穩態格預設＝由下而上 FY30'),
    ('每 MW 電費', 'US$m/MW/年', '—（overlay 預設關閉）', f"{tkv('IF_PowerCost', GEN[0]):.3f}", f"{tkv('IF_PowerCost', GEN[0], '低成本'):.3f}–{tkv('IF_PowerCost', GEN[0], '高成本'):.3f}", 'IF_PowerCost（已含平均用電比與 PUE；與世代無關）', '[Analogy]', '電價、PUE 為 Tokenomics 產業值；CRWV 站點電價找不到（W1 6e）'),
    ('每 MW IT 維護', 'US$m/MW/年', 'overlay 0.15–0.17（預設關閉）', gl('IF_MaintIT'), gr('IF_MaintIT'), 'IF_MaintIT（v5.26；W2 暫代 0.15）', '[Analogy]', '在役世代加權；FY26 0.94 → FY30 1.08'),
    ('每 MW 人員、軟體、水與耗材', 'US$m/MW/年', '—', f"{tkv('IF_StaffSW', GEN[0]):.3f}", f"{tkv('IF_StaffSW', GEN[0], '低成本'):.2f}–{tkv('IF_StaffSW', GEN[0], '高成本'):.2f}", 'IF_StaffSW（v5.26；W2 暫代 0）', '[Analogy]', '全額計入'),
    ('每 MW 財產稅與保險', 'US$m/MW/年', '—', gl('IF_TaxIns'), gr('IF_TaxIns'), 'IF_TaxIns × IF_CapexIT ÷ IF_CapexTotal（v5.26；W2 暫代 0）', '[Analogy]', '只算 CRWV 擁有的 IT 部分（廠房屬房東）；GB300 0.251 × 74.7% ＝ 0.187'),
    ('公司管銷率（扣 D&A、SBC）', '% 營收', '—（含在 EBITDA 率）', f"{(LQ['sm'] - LQ['smSbc'] + LQ['ga'] - LQ['gaSbc']) / LQ['revenue']:.2%}", f"{(LQ['sm'] - LQ['smSbc'] + LQ['ga'] - LQ['gaSbc']) / LQ['revenue']:.2%}–{(LQ['sm'] + LQ['ga']) / LQ['revenue']:.2%}（GAAP）",
     f"10-Q {LQ['periodEnd']}：（銷售行銷 {LQ['sm']} − SBC {LQ['smSbc']} ＋ 一般管理 {LQ['ga']} − SBC {LQ['gaSbc']}）÷ 營收 {LQ['revenue']}", '[Derived]', '各期持平；S&M／G&A 內 D&A 視為 0 [Assumed]'),
    ('期初在役世代組合（6/30，1,500 MW）', '%', '—', '／'.join(f"{GC[k]} {v:.1%}" for k, v in CO['fleet']['openMix']['mix'].items()), 'Hopper 24–31%、GB200 20–42%、GB300 30–53%', 'W1 第 6b 步年份分層（data/permw_inputs_20261007.json）', '[Assumed]', 'MW 錨點 [Verified]；各層占比公司未揭露'),
    ('新增 MW 世代組合', '%', '—', '2H26 GB300 80%／GB200 20%；FY27 GB300 50%／VR200 50%；FY28–30 VR200 100%', 'Rubin Ultra 版（FY29–30）只作敏感度', 'W2 工作單預設（fleet.newMix）', '[Assumed]', 'GB300 為 2026 主力、VR 2026 Q2 首套驗證'),
    ('MW 口徑', '—', '未定義', 'IT 關鍵電力（meta.mwBasis＝IT）', '設施口徑：每 MW 數字 × 1.2', 'W1 第 6a 步：房東以 critical IT 計、GPU 密度接近 IT 口徑', '[Assumed]', '公司未定義 active power 口徑；換算已實作、本次不採（拆解 ⑤ ＝ 0）'),
    ('每 MW 年收入（100% 計費）', 'US$m/MW/年', '11.2、11.5、11.5、11.0、10.5', '同左（未改）', '—', 'defaults.m.revMW（備案 revenue＝legacy）', '[Assumed]（v4.5）', 'GPU 小時價格路線已實作、預設不啟用（拆解 ④ ＝ 0）'),
    ('電力／維護 overlay', '—', '可開啟（預設關閉）', 'bottomUp 下自動停用（公式改為 0）', '—', 'build_xlsx：避免與由下而上重複扣', '—', '舊方法組合維持原行為'),
]
for row in PAR:
    for j, v in enumerate(row): put(ws, r, 1 + j, v, wrap=True, font=BOLD if j == 0 else BLACK)
    r += 1

# ======================= 變動拆解 =======================
ws = sheet('變動拆解', [52, 12, 12, 12, 12, 12, 12, 12, 12, 12])
ws['A1'] = '目標價變動拆解 v4.5 → v4.6（已決定事項 12(3)；0 截斷口徑；US$／股）'; ws['A1'].font = TITLE
ws['A2'] = ('(a) 時間推移＝0（評價日 2026-06-30 不變）、(b) 實際數更新＝0、(c) 假設變更＝0（共用輸入未改）、(d) 方法變更＝總變動。(d) 依 ① → ⑤ 逐項切換（每步只改一項、前一步為起點）；'
            '「單獨切換」＝只改該項、其他維持 v4.5。數值由 scripts/attrib_permw.py 以 Excel（LibreOffice）重算各組合取得。'); ws['A2'].font = SMALL; ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:J2'); ws.row_dimensions[2].height = 40
r = 4
r = hdr(ws, r, ['項目'] + [f'{SCK[s]}｜{h}' for s, _ in SCN for h in ('依序', '單獨', '順序依賴')])
top = r
put(ws, r, 1, 'v4.5 加權目標價', font=BOLD)
for j, (s, _) in enumerate(SCN): put(ws, r, 2 + 3 * j, att[SCK[s]]['v45']['blend'], USD, BLUE)
r += 1
for lab, k in (('(a) 時間推移', 'a'), ('(b) 實際數更新', 'b'), ('(c) 假設變更', 'c')):
    put(ws, r, 1, lab)
    for j, (s, _) in enumerate(SCN): put(ws, r, 2 + 3 * j, att[SCK[s]]['abcd'][k], USD, BLUE)
    r += 1
put(ws, r, 1, '(d) 方法變更（＝下列 ① 至 ⑤ 合計）', font=BOLD); dr = r; r += 1
s0 = r
for k, nm in AT['steps']:
    put(ws, r, 1, '　' + nm)
    for j, (s, _) in enumerate(SCN):
        q, al = att[SCK[s]]['seq'][k]['blend'], att[SCK[s]]['alone'][k]['blend']
        put(ws, r, 2 + 3 * j, q, USD, BLUE); put(ws, r, 3 + 3 * j, al, USD, BLUE); put(ws, r, 4 + 3 * j, f'={CL(2 + 3 * j)}{r}-{CL(3 + 3 * j)}{r}', USD)
    r += 1
for j in range(3):
    put(ws, dr, 2 + 3 * j, f'=SUM({CL(2 + 3 * j)}{s0}:{CL(2 + 3 * j)}{r - 1})', USD)
    put(ws, dr, 3 + 3 * j, f'=SUM({CL(3 + 3 * j)}{s0}:{CL(3 + 3 * j)}{r - 1})', USD)
put(ws, r, 1, 'v4.6 加權目標價（＝v4.5＋(a)＋(b)＋(c)＋(d)）', font=BOLD)
for j in range(3): put(ws, r, 2 + 3 * j, f'={CL(2 + 3 * j)}{top}+SUM({CL(2 + 3 * j)}{top + 1}:{CL(2 + 3 * j)}{top + 3})+{CL(2 + 3 * j)}{dr}', USD)
vr = r; r += 1
put(ws, r, 1, 'v4.6 成品加權目標價（取自成品）')
for j, (s, _) in enumerate(SCN): put(ws, r, 2 + 3 * j, tgt('v46', s), USD, BLUE)
r += 1
put(ws, r, 1, '檢查：拆解合計 − 成品（須 < 0.01）', font=BOLD)
for j in range(3): put(ws, r, 2 + 3 * j, f'=IF(ABS({CL(2 + 3 * j)}{vr}-{CL(2 + 3 * j)}{r - 1})<0.01,"通過","不通過")')
r += 2
r = sec(ws, r, '兩條腿的依序拆解（DCF 腿推到 2027-06-30；EV/EBITDA 腿錨定 FY29、折回 2027-06-30；權重 DCF 45％／EV 55％，DCF 失效時 0／100％）', 10)
r = hdr(ws, r, ['項目'] + [f'{SCK[s]}｜{h}' for s, _ in SCN for h in ('DCF 腿', 'EV/EBITDA 腿', '加權')])
for lab, key in [('v4.5', 'v45')] + [(nm, k) for k, nm in AT['steps']] + [('v4.6', 'v46')]:
    put(ws, r, 1, lab if key in ('v45', 'v46') else '　' + lab, font=BOLD if key in ('v45', 'v46') else BLACK)
    for j, (s, _) in enumerate(SCN):
        o = att[SCK[s]]
        vals = o[key] if key in ('v45', 'v46') else o['seq'][key]
        for m, mm in enumerate(('dcf', 'ev', 'blend')): put(ws, r, 2 + 3 * j + m, vals[mm], USD, BLUE)
    r += 1
ws.cell(row=r, column=1, value='讀法：v4.5、v4.6 列為水準；中間各列為該步的變動。保守、積極情境的「順序依賴」大，是 0 截斷與 DCF 失效（權重歸零）的非線性所致——先改建置成本後 DCF 腿已為 0，營運成本再改只影響 EV/EBITDA 腿。').font = SMALL

# ======================= Tokenomics 參考線 =======================
ws = sheet('Tokenomics參考線', [52, 12] + [11] * 5 + [60])
ws['A1'] = 'Tokenomics 參考線（v4.6 基準；只作對照，不入損益）'; ws['A1'].font = TITLE
r = hdr(ws, 3, ['項目', '單位'] + PER + ['說明'])
REF = [('每 MW 經濟持有成本（不賠錢下限）', 'IF_HoldEcon 在役世代加權（含廠房資本回收；CRWV 的廠房在租金，口徑較寬）'),
       ('每 MW 年收入（模型採用，100% 計費時數）', '模型輸入（v4.5 既有）'), ('每 MW 年收入 ÷ 經濟持有成本', '< 1：計費單價不足以回收含廠房的持有成本'),
       ('每 GPU 小時經濟持有成本（GPU 數加權）', 'IF_GPUhrEcon（100% 時數、不含利潤）'), ('隱含每 GPU 小時價格（反算對照，不是輸入）', '＝每 MW 年收入 ÷（每 MW GPU 數 × 8,760）'),
       ('隱含價格 ÷ GPU 小時持有成本', ''), ('每 MW GPU 數（世代加權）', 'IF_GPUsPerGW ÷ 1000'),
       ('同業｜IREN–Microsoft 每 MW 年收入（GB300，5 年合約）', '[Derived]：$9.7bn ÷ 5 年 ÷ 200 MW critical IT'),
       ('現金成本合計（含租金）', '模型（每MW經濟性彙總）'), ('EBITDA', '模型（每MW經濟性彙總）'), ('D&A（模型車隊折舊）', '模型'),
       ('每 MW 建置成本 ÷ 壽命（在役世代加權；D&A 參考）', 'IF_CapexIT ÷ 6 年（在役世代加權）')]
for k, nt in REF:
    put(ws, r, 1, k); put(ws, r, 2, '倍' if '÷' in k and '壽命' not in k else ('US$/GPU-hr' if 'GPU 小時' in k and '÷' not in k else ('顆/MW' if 'GPU 數' in k else 'US$m/MW')), font=SMALL)
    vals = g('v46', B, PM + k)
    for i in range(5): put(ws, r, 3 + i, vals[i], NUM, BLUE)
    put(ws, r, 8, nt, font=SMALL, wrap=True); r += 1
for k in [x for x in D['v46'][B] if x.startswith(PM + '市場價格｜')]:
    put(ws, r, 1, k[len(PM):]); put(ws, r, 2, 'US$/GPU-hr', font=SMALL); put(ws, r, 3, D['v46'][B][k][0], NUM, BLUE); put(ws, r, 8, '單點（W1 第 6c 步；牌價為隨需、指數為市場現行租期）', font=SMALL); r += 1
r += 1
r = sec(ws, r, 'Tokenomics 各世代值（基準；$B/GW ＝ US$m／IT MW）', 8)
r = hdr(ws, r, ['名稱'] + [GC[x] for x in GEN] + ['', '', '單位'])
for n in ['IF_CapexIT', 'IF_CapexTotal', 'IF_PowerCost', 'IF_MaintIT', 'IF_StaffSW', 'IF_TaxIns', 'IF_DeprLifeIT', 'IF_HoldEcon', 'IF_GPUhrEcon', 'IF_GPUsPerGW']:
    put(ws, r, 1, f"{n}（{TK['items'][n]['label']}）")
    for j, x in enumerate(GEN): put(ws, r, 2 + j, tkv(n, x) if x in TK['items'][n]['values'] else None, '#,##0.000', BLUE)
    put(ws, r, 8, TK['items'][n]['unit'], font=SMALL); r += 1

# ======================= 已知限制 =======================
ws = sheet('已知限制', [40, 70, 50])
ws['A1'] = '已知限制（v4.6）'; ws['A1'].font = TITLE
r = hdr(ws, 3, ['限制', '內容', '對結果的影響方向'])
LIM = [
    ('收入採備案（每 MW 年收入為輸入）', 'GB200／GB300／VR200 長約 GPU 小時價格不足兩個獨立來源，VR200 完全找不到；GPU 小時路線已實作未啟用', '方向不定；模型每 MW 收入 10.0 高於 Q2 年化 8.2，若以 Q2 校準收入會下修、目標價下降'),
    ('續約價格衰退未做', '合約到期後的續約單價下降未建模（範圍外：CRWV 原待辦）', '目標價偏高（FY29–30 收入高估）'),
    ('Tokenomics 壽命期價格係數 L 未接', 'Tokenomics 的價格隨時間衰減係數未引用', '目標價偏高'),
    ('管銷率各期持平 5.51%', '未反映規模經濟', '目標價偏低（若管銷率隨規模下降）；GAAP 9.24% 則基準 $21.25（−$8.43）'),
    ('期初世代組合為估計', '公司未揭露機隊世代；年份分層 [Assumed]', '影響每 MW IT 維護、稅險與參考線；Rubin Ultra 版敏感度基準 $22.09'),
    ('Tokenomics 營運成本為滿載穩態口徑', 'Q2 實際營運成本（租金前、不含管銷）0.88／MW，模型 FY26 2.09（保固期、爬坡 MW、電費轉嫁可能使實際偏低）', '若實際較低的狀態延續：EBITDA 率上升、目標價上升（Tokenomics 低成本敏感度基準 $72.45）'),
    ('租金路徑低於 Q2 實際', '模型每 MW 年租金 1.60（FY26）對 Q2 2.08；屬 v4.5 既有（具名站點基準未做 MW 驅動）', '目標價偏高'),
    ('MW 口徑未證實', 'active power 是 IT 或設施口徑公司未定義；預設 IT', '設施口徑時每 MW 數字 × 1.2，利潤率不變'),
    ('車隊折舊高於實際（待辦 6）', '模型 D&A 高於 Q2 實際走勢（含 CIP）', 'EBIT 偏低；EBITDA 與目標價（EV/EBITDA 腿）不受影響'),
    ('LibreOffice 驗證限制', '情境區間運算表在 LibreOffice 重算非基準情境時會殘留中間值；verify 已改以凍結運算表輸出的副本比對（成品 Excel 由真正的 Excel 開啟時不受影響）', '無數字影響；需 Andy 以真正的 Excel 開啟確認'),
]
for row in LIM:
    for j, v in enumerate(row): put(ws, r, 1 + j, v, wrap=True, font=BOLD if j == 0 else BLACK)
    r += 1
wb.save(OUT)
if SUMJ:
    json.dump({'conclusion': CONC, 'summary': SUM, 'q2': QS, 'qmw': QMW}, open(SUMJ, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('saved', OUT)
