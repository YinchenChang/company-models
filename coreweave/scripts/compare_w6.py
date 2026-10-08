# W6（2026-10-09）：v4.7 → v4.8（接 Tokenomics v5.31）對照報告（工作單 W6 第 6 步）。
# 讀兩版 Excel（LibreOffice 開啟、三情境切換「情境選擇」）的「每MW經濟性」彙總列、IT 維護機齡列、加權目標價、評等、融資缺口與 EBITDA 率，
# 兩版「公司實況驗證」頁（參數驗證、Q2 逐項對帳；基準情境快取值），Tokenomics 快照前後值（v5.27 自 git 歷史），
# attrib_w6.py（變動拆解）與兩版 permw_sens.json（敏感度快照；v4.7 自 git 歷史），寫成對照 Excel 與 JSON（供 md 報告）。
# 用法：python3 scripts/compare_w6.py --v47 v4.7.xlsx --v48 v4.8.xlsx --attrib attrib_w6.json --out 報告.xlsx [--json 資料.json]
import sys, os, json, argparse, subprocess
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
from uno_q import Workbook
from openpyxl import Workbook as XW
from openpyxl.styles import Font, PatternFill
from compare_w5 import SUM, SCN, CODE, P, cvtab

REF = os.environ.get('V47_REF', '8671a11')
AGE = ['保固期內占比（平均在役 MW）', '每 MW IT 維護（世代加權）', '每 MW IT 維護｜保固期內部分（IF_MaintITWarr）', '每 MW IT 維護｜保固期滿部分（IF_MaintITPost）',
       '每 MW IT 維護｜等值費率 IF_MaintIT（對照，不入損益）', '每 MW 經濟持有成本（不賠錢下限）', '定價倍數 k（既有占比 × k_既有 ＋（1 − 既有占比）× k_新約）']


def grab(x):
    out = {}
    with Workbook(x) as wb:
        sel = wb.cell('輸入與假設', '情境選擇', prefix=True); s0 = wb.get(sel)
        def row(sh, nm):
            try: return [wb.get(wb.cell(sh, nm, col=c)) for c in 'CDEFG']
            except KeyError: return None
        for s in (1, 2, 3):
            wb.set(sel, s); o = {nm: row('每MW經濟性', nm) for nm in SUM + AGE}
            o['tp'] = wb.get(wb.cell('評價_DCF與目標價', '加權目標價'))
            o['rating'] = CODE[round(wb.get(wb.cell('評價_DCF與目標價', '目標價區間｜點位評等代碼（1＝買進、0＝中立、−1＝賣出）')))]
            o['gap'] = max(0.0, -wb.get(wb.cell('各期收支', '融資前累積現金', col='G')))
            o['ebm'] = row('運營_產能與收入', 'EBITDA 率（損益與資金共用）')
            out[s] = o
        wb.set(sel, s0)
    return out


def git_json(path):
    return json.loads(subprocess.run(['git', '-C', ROOT, 'show', f'{REF}:coreweave/{path}'], capture_output=True, text=True, check=True).stdout)


def tkdiff():
    a = git_json('data/tokenomics_snapshot_v5.27.json')['items']
    co = json.load(open(os.path.join(ROOT, 'company.json'), encoding='utf-8'))
    b = json.load(open(os.path.join(ROOT, co['tokenomics']['snapshotFile']), encoding='utf-8'))['items']
    rows = []
    for n, y in b.items():
        x = a.get(n)
        if y['kind'] == 'gen_cost':
            for g, v in y['values'].items():
                for c in ('低成本', '基準', '高成本'):
                    u = x['values'][g][c] if x else None
                    rows.append([n, g, c, u, v[c], (v[c] / u - 1) if isinstance(u, (int, float)) and u else None])
        else:
            u = x['values'] if x else None
            rows.append([n, '—', '基準', u, y['values'], (y['values'] / u - 1) if isinstance(u, (int, float)) and u else None])
    return rows


def main():
    ap = argparse.ArgumentParser()
    for k in ('v47', 'v48', 'attrib', 'out'): ap.add_argument('--' + k, required=True)
    ap.add_argument('--json'); a = ap.parse_args()
    A, B = grab(os.path.abspath(a.v47)), grab(os.path.abspath(a.v48))
    A[2]['cv'], B[2]['cv'] = cvtab(os.path.abspath(a.v47)), cvtab(os.path.abspath(a.v48))
    AT = json.load(open(a.attrib, encoding='utf-8'))
    SN = {'v47': git_json('permw_sens.json'), 'v48': json.load(open(os.path.join(ROOT, 'permw_sens.json'), encoding='utf-8'))}
    TK = tkdiff()
    data = {'v47': A, 'v48': B, 'attrib': AT, 'sens': SN, 'tk': TK}
    H, T, S_ = Font(bold=True, color='FFFFFF'), Font(bold=True, size=13), Font(size=9, color='555555')
    FH = PatternFill('solid', fgColor='1F3864')
    wb = XW(); ws = wb.active; ws.title = '摘要'
    def hdr(ws, r, xs):
        for j, x in enumerate(xs):
            c = ws.cell(row=r, column=1 + j, value=x); c.font = H; c.fill = FH
    ws['A1'] = 'CoreWeave v4.7 → v4.8（接 Tokenomics v5.31：GB300 機架價格、IT 維護機齡兩段）'; ws['A1'].font = T
    ws['A2'] = 'k 隨錨重算（價格不變，k＝證據價格 ÷ v5.31 同世代成本）；2023 年底 MW 依 10-K 原文 70；MW 口徑查證後維持 IT（10-K、10-Q 無原文定義）'; ws['A2'].font = S_
    r = 4; hdr(ws, r, ['情境', 'v4.7 目標價', 'v4.8 目標價', '變動', 'v4.7 評等', 'v4.8 評等', 'v4.7 融資缺口', 'v4.8 融資缺口', '缺口變動']); r += 1
    for s in (1, 2, 3):
        for j, v in enumerate([SCN[s], A[s]['tp'], B[s]['tp'], f'=C{r}-B{r}', A[s]['rating'], B[s]['rating'], A[s]['gap'], B[s]['gap'], f'=H{r}-G{r}']):
            ws.cell(row=r, column=1 + j, value=v).number_format = '0.00'
        r += 1
    w = wb.create_sheet('每MW_前後'); w['A1'] = '每 MW（US$m／平均在役 MW／年）：v4.7 → v4.8，三情境'; w['A1'].font = T
    r = 3; hdr(w, r, ['情境', '項目', '版本'] + P); r += 1
    for s in (1, 2, 3):
        for nm in SUM + ['EBITDA 率（模型採用）']:
            for tag, X in (('v4.7', A), ('v4.8', B)):
                v = X[s]['ebm'] if nm.startswith('EBITDA 率') else X[s][nm]
                w.cell(row=r, column=1, value=SCN[s]); w.cell(row=r, column=2, value=nm); w.cell(row=r, column=3, value=tag)
                for i, x in enumerate(v or []): w.cell(row=r, column=4 + i, value=x).number_format = '0.0%' if nm.startswith('EBITDA 率') else '0.000'
                r += 1
    w = wb.create_sheet('IT維護機齡'); w['A1'] = 'IT 維護依機齡兩段（v4.8，三情境；v4.7 為等值費率 IF_MaintIT v5.27）'; w['A1'].font = T
    r = 3; hdr(w, r, ['情境', '項目', '版本'] + P); r += 1
    for s in (1, 2, 3):
        w.cell(row=r, column=1, value=SCN[s]); w.cell(row=r, column=2, value='每 MW IT 維護（世代加權）'); w.cell(row=r, column=3, value='v4.7')
        for i, x in enumerate(A[s]['每 MW IT 維護（世代加權）'] or []): w.cell(row=r, column=4 + i, value=x).number_format = '0.000'
        r += 1
        for nm in AGE:
            w.cell(row=r, column=1, value=SCN[s]); w.cell(row=r, column=2, value=nm); w.cell(row=r, column=3, value='v4.8')
            for i, x in enumerate(B[s][nm] or []): w.cell(row=r, column=4 + i, value=x).number_format = '0.0%' if '占比' in nm else '0.000'
            r += 1
    for tag, X in (('v4.7', A), ('v4.8', B)):
        w = wb.create_sheet(f'公司實況驗證_{tag}'); w['A1'] = f'公司實況驗證（{tag} Excel「公司實況驗證」頁讀值；基準情境）'; w['A1'].font = T
        r = 3; hdr(w, r, ['參數', '單位', 'Tokenomics 值', 'CRWV 實際', '差距', '採用值', '規則', 'Tokenomics 名稱', 'CRWV 實際與標記', '差距原因（機制與證據）', '公司調整', '證據']); r += 1
        for x in X[2]['cv']:
            for j, v in enumerate(x): w.cell(row=r, column=1 + j, value=v)
            r += 1
    w = wb.create_sheet('Tokenomics前後'); w['A1'] = 'Tokenomics 引用名稱前後值（v5.27 → v5.31；$B/GW＝US$m/MW）'; w['A1'].font = T
    r = 3; hdr(w, r, ['名稱', '世代', '成本情境', 'v5.27', 'v5.31', '變動']); r += 1
    for x in TK:
        for j, v in enumerate(x):
            c = w.cell(row=r, column=1 + j, value=v)
            if j == 5 and v is not None: c.number_format = '0.0%'
        r += 1
    w = wb.create_sheet('敏感度'); w['A1'] = '敏感度（建置時快照 permw_sens.json；加權目標價 US$／融資缺口 US$bn；v4.7 → v4.8）'; w['A1'].font = T
    r = 3; hdr(w, r, ['設定', '版本', '保守 目標價', '保守 缺口', '基準 目標價', '基準 缺口', '積極 目標價', '積極 缺口']); r += 1
    names = dict(SN['v47']['cases']); names.update(dict(SN['v48']['cases']))
    for k in dict.fromkeys([c for c, _ in SN['v48']['cases']] + [c for c, _ in SN['v47']['cases']]):
        for tag in ('v47', 'v48'):
            w.cell(row=r, column=1, value=dict(SN[tag]['cases']).get(k, names[k])); w.cell(row=r, column=2, value='v4.7' if tag == 'v47' else 'v4.8')
            for j, sk in enumerate(('low', 'base', 'high')):
                v = SN[tag]['scenarios'][sk].get(k)
                w.cell(row=r, column=3 + 2 * j, value='不適用' if v is None else v[0]); w.cell(row=r, column=4 + 2 * j, value='不適用' if v is None else v[1])
            r += 1
    w = wb.create_sheet('變動拆解'); w['A1'] = '目標價變動拆解（v4.7 → v4.8；0 截斷口徑，US$／股；scripts/attrib_w6.py）'; w['A1'].font = T
    r = 3; hdr(w, r, ['逐項'] + list(AT['scenarios'])); r += 1
    top = r
    for k, nm in AT['steps']:
        w.cell(row=r, column=1, value=nm)
        for j, sc in enumerate(AT['scenarios']): w.cell(row=r, column=2 + j, value=AT['scenarios'][sc]['seq'][k]['blend']).number_format = '+0.00;-0.00'
        r += 1
    w.cell(row=r, column=1, value='合計（＝總變動）')
    for j in range(3): w.cell(row=r, column=2 + j, value=f'=SUM({chr(66 + j)}{top}:{chr(66 + j)}{r - 1})').number_format = '+0.00;-0.00'
    r += 1; w.cell(row=r, column=1, value='總變動（v4.8 − v4.7，讀值）')
    for j, sc in enumerate(AT['scenarios']): w.cell(row=r, column=2 + j, value=AT['scenarios'][sc]['total']).number_format = '+0.00;-0.00'
    r += 1; w.cell(row=r, column=1, value='檢查（合計 − 總變動，應為 0）')
    for j in range(3): w.cell(row=r, column=2 + j, value=f'=IF(ABS({chr(66 + j)}{r - 2}-{chr(66 + j)}{r - 1})<0.01,"通過","不通過")')
    r += 2
    for k, nm in (('a', '(a) 時間推移'), ('b', '(b) 實際數更新（2023 年底 MW）'), ('c', '(c) 假設變更（①＋②a＋②b）'), ('d', '(d) 方法變更（③＋④）')):
        w.cell(row=r, column=1, value=nm)
        for j, sc in enumerate(AT['scenarios']): w.cell(row=r, column=2 + j, value=AT['scenarios'][sc]['abcd'][k]).number_format = '+0.00;-0.00'
        r += 1
    wb.save(a.out); print('saved', a.out)
    if a.json: json.dump(data, open(a.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)


if __name__ == '__main__':
    sys.exit(main())
