# W4（2026-10-08）：v4.6 → v4.7 收入錨定對照報告（工作單 W4 第 6 步）。
# 讀兩版 Excel（LibreOffice 開啟、三情境切換「情境選擇」）的「每MW經濟性」逐列、加權目標價、評等與融資缺口，
# 加上 attrib_w4.py（變動拆解）、q2_check_w4.py（Q2 驗證）、permw_sens.json（敏感度快照）與 company.json 證據表，
# 寫成對照 Excel（數值為讀取值；檢查格為活公式）與 JSON（供 md 報告）。
# 用法：python3 scripts/compare_w4.py --v46 v4.6.xlsx --v47 v4.7.xlsx --attrib attrib_w4.json --q2 q2.json --out 報告.xlsx [--json 資料.json]
import sys, os, json, argparse
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
from uno_q import Workbook
from openpyxl import Workbook as XW
from openpyxl.styles import Font, PatternFill, Alignment

ROWS = ['每 MW 年收入（算力＋服務）', '　其中：算力收入', '現金成本合計（含租金）', 'EBITDA', '稅前', '平均在役 MW 合計',
        '每 MW 年收入（模型採用，100% 計費時數）', '每 MW 年收入（計費後＝× 利用率）', '每 MW 經濟持有成本（不賠錢下限）', '每 MW 年收入 ÷ 經濟持有成本']
W4ROWS = ['錨｜每 MW 經濟持有成本（IF_HoldEcon，在役世代加權）', '錨｜排程 RPO（模型期）', '錨｜平均計費 MW', '長約價產能收入（平均計費 MW × 錨 × k_長約 × 利用率 × 期間長度）',
          'RPO 涵蓋的產能 ÷ 在役計費產能（對照，不驅動）', 'k_長約（依目前成本情境重算）', '隨需占比（輸入）', '定價倍數 k（隨需占比 × k_現貨 ＋（1 − 隨需占比）× k_長約）',
          'Tokenomics 錨 × k：每 MW 年收入（100% 計費時數）', '錨 × k 相對 v4.6 舊輸入 m.revMW',
          '上限檢查｜客戶每 MW 付費 token 營收（IF_RevGWFleet，在役世代加權）', '上限檢查｜CRWV 每 MW 計費收入 ÷ 客戶付費 token 營收']
COST = ['電費', 'IT 維護', '人員、軟體、水與耗材', '財產稅與保險', '租金']
SCN = {1: '保守 4.2 GW', 2: '基準 5.6 GW', 3: '積極 8 GW'}
CODE = {1: '買進', 0: '中立', -1: '賣出'}


def grab(x, w4):
    out = {}
    with Workbook(x) as wb:
        sel = wb.cell('輸入與假設', '情境選擇', prefix=True); s0 = wb.get(sel)
        for s in (1, 2, 3):
            wb.set(sel, s); o = {}
            for nm in ROWS + COST + (W4ROWS if w4 else []):
                o[nm] = [wb.get(wb.cell('每MW經濟性', nm, col=c)) for c in 'CDEFG']
            o['tp'] = wb.get(wb.cell('評價_DCF與目標價', '加權目標價'))
            o['rating'] = CODE[round(wb.get(wb.cell('評價_DCF與目標價', '目標價區間｜點位評等代碼（1＝買進、0＝中立、−1＝賣出）')))]
            o['gap'] = max(0.0, -wb.get(wb.cell('各期收支', '融資前累積現金', col='G')))
            o['ebm'] = [wb.get(wb.cell('運營_產能與收入', 'EBITDA 率（損益與資金共用）', col=c)) for c in 'CDEFG']
            out[s] = o
        wb.set(sel, s0)
    return out


def main():
    ap = argparse.ArgumentParser()
    for k in ('v46', 'v47', 'attrib', 'q2', 'out'): ap.add_argument('--' + k, required=True)
    ap.add_argument('--json'); a = ap.parse_args()
    A, B = grab(os.path.abspath(a.v46), False), grab(os.path.abspath(a.v47), True)
    AT, Q2 = json.load(open(a.attrib, encoding='utf-8')), json.load(open(a.q2, encoding='utf-8'))
    SN = json.load(open(os.path.join(ROOT, 'permw_sens.json'), encoding='utf-8'))
    CO = json.load(open(os.path.join(ROOT, 'company.json'), encoding='utf-8')); AM = CO['pricing']['anchorMultiple']
    TK = json.load(open(os.path.join(ROOT, CO['tokenomics']['snapshotFile']), encoding='utf-8'))
    tk = lambda n, g: TK['items'][n]['values'][g]['基準']
    data = {'v46': A, 'v47': B, 'attrib': AT, 'q2': Q2, 'sens': SN}
    P = ['FY26（2H）', 'FY27', 'FY28', 'FY29', 'FY30']
    H, T, S_ = Font(bold=True, color='FFFFFF'), Font(bold=True, size=13), Font(size=9, color='555555')
    FH = PatternFill('solid', fgColor='1F3864')
    wb = XW(); ws = wb.active; ws.title = '摘要'
    def hdr(ws, r, xs):
        for j, x in enumerate(xs):
            c = ws.cell(row=r, column=1 + j, value=x); c.font = H; c.fill = FH
    b6, b7 = A[2], B[2]
    ws['A1'] = 'CoreWeave v4.6 → v4.7 收入錨定對照（每 MW 收入＝Tokenomics IF_HoldEcon × 定價倍數 k）'; ws['A1'].font = T
    ws['A2'] = f"Tokenomics {CO['tokenomics']['version']}（commit {CO['tokenomics']['commit'][:7]}）；k_長約 {AM['long']['base']}（{AM['long']['low']}–{AM['long']['high']}）、k_現貨 {AM['spot']['base']}（{AM['spot']['low']}–{AM['spot']['high']}）；k＝隨需占比 × k_現貨＋（1 − 隨需占比）× k_長約，隨需占比 {AM['onDemandShare']['base']:.0%}"
    ws['A2'].font = S_
    r = 4; hdr(ws, r, ['情境', 'v4.6 加權目標價', 'v4.7 加權目標價', '變動', 'v4.6 評等', 'v4.7 評等', 'v4.6 融資缺口', 'v4.7 融資缺口', '變動']); r += 1
    for s in (1, 2, 3):
        for j, v in enumerate([SCN[s], A[s]['tp'], B[s]['tp'], f'=C{r}-B{r}', A[s]['rating'], B[s]['rating'], A[s]['gap'], B[s]['gap'], f'=H{r}-G{r}']):
            c = ws.cell(row=r, column=1 + j, value=v); c.number_format = '0.00'
        r += 1
    r += 1; hdr(ws, r, ['基準：每 MW（US$m／MW／年）'] + P); r += 1
    for lab, src, nm in [('v4.6 每 MW 年收入（100% 計費）', A, '每 MW 年收入（模型採用，100% 計費時數）'), ('v4.7 錨（IF_HoldEcon 在役加權）', B, 'Tokenomics 錨 × k：每 MW 年收入（100% 計費時數）'),
                         ('v4.7 定價倍數 k', B, '定價倍數 k（隨需占比 × k_現貨 ＋（1 − 隨需占比）× k_長約）'), ('v4.7 RPO 覆蓋率（對照，不驅動）', B, 'RPO 涵蓋的產能 ÷ 在役計費產能（對照，不驅動）')]:
        vals = src[2][nm] if lab != 'v4.7 錨（IF_HoldEcon 在役加權）' else src[2]['錨｜每 MW 經濟持有成本（IF_HoldEcon，在役世代加權）']
        ws.cell(row=r, column=1, value=lab)
        for i, v in enumerate(vals): ws.cell(row=r, column=2 + i, value=v).number_format = '0.000'
        r += 1
    ws.cell(row=r, column=1, value='v4.7 每 MW 年收入（100% 計費＝錨 × k）')
    for i, v in enumerate(b7['每 MW 年收入（模型採用，100% 計費時數）']): ws.cell(row=r, column=2 + i, value=v).number_format = '0.000'
    r += 1
    ws.cell(row=r, column=1, value='檢查：錨 × k − 每 MW 年收入（應為 0）')
    for i in range(5): ws.cell(row=r, column=2 + i, value=f'={chr(66 + i)}{r - 4}*{chr(66 + i)}{r - 3}-{chr(66 + i)}{r - 1}').number_format = '0.000000'
    r += 2
    ws.column_dimensions['A'].width = 44
    for c in 'BCDEFGHI': ws.column_dimensions[c].width = 14
    # 每 MW 前後
    w2 = wb.create_sheet('每MW_前後'); r = 1
    for s in (1, 2, 3):
        hdr(w2, r, [SCN[s], '版本'] + P); r += 1
        for nm in ROWS + COST:
            for tag, src in (('v4.6', A), ('v4.7', B)):
                w2.cell(row=r, column=1, value=nm); w2.cell(row=r, column=2, value=tag)
                for i, v in enumerate(src[s][nm]): w2.cell(row=r, column=3 + i, value=v).number_format = '0.000'
                r += 1
            w2.cell(row=r, column=1, value=nm); w2.cell(row=r, column=2, value='差異')
            for i in range(5): w2.cell(row=r, column=3 + i, value=f'={chr(67 + i)}{r - 1}-{chr(67 + i)}{r - 2}').number_format = '0.000'
            r += 1
        for tag, src in (('v4.6', A), ('v4.7', B)):
            w2.cell(row=r, column=1, value='EBITDA 率'); w2.cell(row=r, column=2, value=tag)
            for i, v in enumerate(src[s]['ebm']): w2.cell(row=r, column=3 + i, value=v).number_format = '0.0%'
            r += 1
        r += 1
    w2.column_dimensions['A'].width = 52
    # 錨與 k（三情境）
    w3 = wb.create_sheet('錨與k'); r = 1
    for s in (1, 2, 3):
        hdr(w3, r, [SCN[s]] + P); r += 1
        for nm in W4ROWS:
            w3.cell(row=r, column=1, value=nm)
            for i, v in enumerate(B[s][nm]): w3.cell(row=r, column=2 + i, value=v).number_format = '0.000'
            r += 1
        r += 1
    w3.column_dimensions['A'].width = 70
    # 證據表
    w4 = wb.create_sheet('證據表'); hdr(w4, 1, ['證據', '世代', '價格', '單位', 'Tokenomics 名稱', 'Tokenomics 基準值', '倍數（公式）', '用途', '合約型態', '期間', '來源', '文件日期', '標記', '網址', '說明'])
    for j, e in enumerate(AM['evidence'], start=2):
        for k, v in enumerate([e['label'], e['gen'], e['price'], e['unit'], e['tkName'], tk(e['tkName'], e['gen']), f'=C{j}/F{j}', e['use'], e['contract'], e['term'], e['source'], e['date'], e['tag'], e['url'], e['note']]):
            w4.cell(row=j, column=1 + k, value=v)
    j += 2
    for cm in AM['contractMix']: w4.cell(row=j, column=1, value='合約組合｜' + cm['label']); w4.cell(row=j, column=3, value=cm['value'] if cm['value'] is not None else cm['tag']); w4.cell(row=j, column=11, value=cm['source']); j += 1
    for nf in AM['notFound']: w4.cell(row=j, column=1, value='找不到'); w4.cell(row=j, column=11, value=nf); j += 1
    w4.column_dimensions['A'].width = 44; w4.column_dimensions['K'].width = 60
    # Q2 驗證
    w5 = wb.create_sheet('Q2驗證'); hdr(w5, 1, ['項目', '模型首期', 'Q2', '貢獻（US$m/MW/年）'])
    keys = ['B', 'U', 'K', 'A', 'O']
    for j, ((nm, v), k) in enumerate(zip(Q2['steps'], keys), start=2):
        w5.cell(row=j, column=1, value=nm); w5.cell(row=j, column=2, value=Q2['model'][k]); w5.cell(row=j, column=3, value=Q2['q2'][k]); w5.cell(row=j, column=4, value=v)
    j += 1; w5.cell(row=j, column=1, value='合計（公式）'); w5.cell(row=j, column=4, value=f'=SUM(D2:D{j - 1})')
    w5.cell(row=j + 1, column=1, value='總差距（模型 − Q2）'); w5.cell(row=j + 1, column=2, value=Q2['model']['R']); w5.cell(row=j + 1, column=3, value=Q2['q2']['R']); w5.cell(row=j + 1, column=4, value=f'=B{j + 1}-C{j + 1}')
    w5.cell(row=j + 2, column=1, value='檢查（|合計 − 總差距| < 0.01）'); w5.cell(row=j + 2, column=4, value=f'=IF(ABS(D{j}-D{j + 1})<0.01,"通過","不通過")')
    w5.cell(row=j + 3, column=1, value='Q2 隱含 k（未調整：年化營收 ÷ 在役 MW ÷ 錨）'); w5.cell(row=j + 3, column=3, value=Q2['q2']['kRaw'])
    w5.column_dimensions['A'].width = 50
    # 敏感度
    w6 = wb.create_sheet('敏感度'); hdr(w6, 1, ['設定'] + [f'{SCN[s]} 目標價' for s in (1, 2, 3)] + [f'{SCN[s]} 融資缺口' for s in (1, 2, 3)])
    for j, (k, nm) in enumerate(SN['cases'], start=2):
        w6.cell(row=j, column=1, value=nm)
        for i, sk in enumerate(('low', 'base', 'high')):
            v = SN['scenarios'][sk][k]
            w6.cell(row=j, column=2 + i, value='不適用' if v is None else v[0]); w6.cell(row=j, column=5 + i, value='不適用' if v is None else v[1])
    w6.column_dimensions['A'].width = 36
    # 變動拆解
    w7 = wb.create_sheet('變動拆解'); hdr(w7, 1, ['情境', '項目', '加權目標價', 'DCF 腿', 'EV/EBITDA 腿', '融資缺口']); r = 2
    for s in (1, 2, 3):
        o = AT['scenarios'][{1: '保守', 2: '基準', 3: '積極'}[s]]; top = r
        for lab in ('(a) 時間推移', '(b) 實際數更新', '(c) 假設變更'):
            w7.cell(row=r, column=1, value=SCN[s]); w7.cell(row=r, column=2, value=lab); w7.cell(row=r, column=3, value=0.0); r += 1
        for k, nm in AT['steps']:
            q = o['seq'][k]
            for j, v in enumerate([SCN[s], '(d) ' + nm, q['blend'], q['dcf'], q['ev'], q['gap']]): w7.cell(row=r, column=1 + j, value=v)
            r += 1
        w7.cell(row=r, column=1, value=SCN[s]); w7.cell(row=r, column=2, value='合計（公式）'); w7.cell(row=r, column=3, value=f'=SUM(C{top}:C{r - 1})'); r += 1
        w7.cell(row=r, column=1, value=SCN[s]); w7.cell(row=r, column=2, value='v4.7 − v4.6'); w7.cell(row=r, column=3, value=o['v47']['blend'] - o['v46']['blend']); r += 1
        w7.cell(row=r, column=1, value=SCN[s]); w7.cell(row=r, column=2, value='檢查（誤差 < 0.01）'); w7.cell(row=r, column=3, value=f'=IF(ABS(C{r - 2}-C{r - 1})<0.01,"通過","不通過")'); r += 2
    w7.column_dimensions['B'].width = 56
    wb.save(a.out); print('saved', a.out)
    if a.json: json.dump(data, open(a.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)


if __name__ == '__main__':
    main()
