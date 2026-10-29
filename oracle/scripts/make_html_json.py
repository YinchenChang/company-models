#!/usr/bin/env python3
"""產生 oracle/html.json（tools/company_html 的設定檔）：列號取自 build_oracle.py 的列號對照（ORCL_ROWMAP），
In 頁列號取自成品 xlsx 的鍵欄。build_oracle.py 改版後重跑本檔，再跑 tools/company_html/build_company_html.py oracle。

用法：ORCL_ROWMAP=rowmap.json python3 oracle/build_oracle.py raw.xlsx ；python3 oracle/scripts/make_html_json.py rowmap.json
"""
import glob, json, os, sys
import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = json.load(open(sys.argv[1], encoding="utf-8"))
X = sorted(glob.glob(os.path.join(HERE, "dist", "*.xlsx")))[0]
wb = openpyxl.load_workbook(X, data_only=True)
INR = {wb["In"].cell(r, 1).value: r for r in range(4, wb["In"].max_row + 1) if wb["In"].cell(r, 1).value}
SCN_H = 1 + next(r for r in range(4, wb["In"].max_row + 1) if "情境驅動對照" in str(wb["In"].cell(r, 1).value or ""))
TR, PR, CR, VAL = R["TR"], R["PR"], R["CR"], R["VAL"]
DC, SO, CON = VAL["DC"], VAL["SO"], VAL["CON"]
TB, PB, CB = TR["B"], PR["B"], CR["B"]
G = R["GATE"]


def drv(key, fmt, out_lab, out_ref):
    r = INR[key]
    return [f"@In!B{r}", [f"@In!E{r}|{fmt}", " ", f"@In!C{r}"], f"@In!G{r}|tag", [out_lab, out_ref]]


def met(name, key, fmt, chart=False, src="PL", cols="D{r}:I{r}", rows=None):
    rows = rows or {s: f"@{src}!" + cols.format(r=PR[s][key]) for s in "ABC"}
    m = {"name": name, "fmt": fmt, "rows": rows}
    if chart: m["chart"] = True
    return m


cfg = {
    "company": "Oracle",
    "ticker": "ORCL",
    "thesis": "@README!A2",
    "question": "FY2029 的 OCI 營收與營業利益率，相對於資本支出與租約的擴張，以及資金缺口（自由現金流為負）要如何補",
    "scenarios": [{"key": "A", "name": "情境 A（樂觀）"}, {"key": "B", "name": "情境 B（基準）"}, {"key": "C", "name": "情境 C（保守）"}],
    "sections": [
        {"id": "thesis", "title": "命題與核心問題", "blocks": [
            {"type": "kv", "items": [["@Val!A" + str(CON["rec"]), "@Val!C" + str(CON["rec"])], ["@Val!A" + str(CON["tp"]), f"@Val!C{CON['tp']}|n2"],
                                     ["@Val!A" + str(CON["px"]), f"@Val!C{CON['px']}|n2"], ["@Val!A" + str(CON["up"]), f"@Val!C{CON['up']}|sp1"],
                                     ["@CF!A" + str(G), [f"@CF!C{G}|n1", "／", f"@CF!D{G}|n1", "／", f"@CF!E{G}|n1"]]]},
            {"type": "h3", "text": "重點觀察"},
            {"type": "list", "items": ["@README!A5", "@README!A6", "@README!A7"]},
            {"type": "table", "caption": f"@Thesis!A{R['SUM'][0]-1}", "grid": True,
             "head": [f"@Thesis!A{R['SUM'][0]}", f"@Thesis!B{R['SUM'][0]}", f"@Thesis!C{R['SUM'][0]}", f"@Thesis!D{R['SUM'][0]}", f"@Thesis!E{R['SUM'][0]}"],
             "rowgen": {"sheet": "Thesis", "rows": f"{R['SUM'][0]+1}-{R['SUM'][1]}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "", "n1", "n1", "n1"],
                        "row_fmt": {str(R['SUM'][0] + k): {"C": "p1", "D": "p1", "E": "p1"} for k in (3, 4, 7)}}},
            {"type": "kv", "items": [["會計年度與口徑", "@README!A9"], ["營收分部層", "@README!A10"]]}]},
        {"id": "drivers", "title": "驅動因子 → 推導量", "blocks": [
            {"type": "table", "caption": "情境驅動（三情境並排，進加權目標價；In 頁情境驅動對照）", "head": [f"@In!{c}{SCN_H}" for c in "BDEFCG"],
             "rows": [[f"@In!B{INR[k]}", f"@In!D{INR[k]}|{f}", f"@In!E{INR[k]}|{f}", f"@In!F{INR[k]}|{f}", f"@In!C{INR[k]}", f"@In!G{INR[k]}|tag"]
                      for k, f in (("scn_pace", "n0"), ("scn_cap", "n0"), ("scn_delay", "n0"), ("scn_w", "p0"))]},
            {"type": "table", "caption": "驅動 → 推導量（情境 B 基準值；敏感度低／高只用於敏感度頁、不進目標價）", "head": ["輸入", "基準值", "標記", "推導量（情境 B）"], "rows": [
                drv("pace_B", "n0", "已連網 AI MW（FY2029 年底）：", f"@Thesis!G{TB['conn']}|n0"),
                drv("delay_B", "n0", "計費 MW（FY2029）：", f"@Thesis!G{TB['bill']}|n0"),
                drv("k_mix", "n3", "每 MW 年收入（FY2029）：", f"@Thesis!G{TB['pm']}|n2"),
                drv("ox_vr200", "n2", "AI 營運費用（FY2029）：", f"@Thesis!G{TB['opex']}|n1"),
                drv("lease_unc", "n0", "AI 機房租金（FY2029）：", f"@Thesis!G{TB['rent']}|n1"),
                drv("capex_g27", "n1", "AI 資本支出（FY2028）：", f"@Thesis!F{TB['cx']}|n1"),
                drv("life_ai", "n0", "AI 折舊（FY2029）：", f"@Thesis!G{TB['dep']}|n1"),
                drv("m_sw", "p1", "雲端與軟體分部利潤（FY2029）：", f"@PL!F{PB['mcs']}|n1"),
                drv("gLT_saas", "p1", "SaaS 營收（FY2032）：", f"@PL!I{PB['saas']}|n1"),
                drv("prepay_cov", "p1", "客戶預付流入（FY2028）：", f"@CF!E{CB['ppin']}|n1"),
            ]},
            {"type": "table", "caption": "命題模組：OCI AI 算力（情境 B）", "grid": True,
             "head": ["@Thesis!A3", "@Thesis!B3", "@Thesis!C3", "@Thesis!D3", "@Thesis!E3", "@Thesis!F3", "@Thesis!G3", "@Thesis!H3", "@Thesis!I3", "@Thesis!J3"],
             "rowgen": {"sheet": "Thesis", "rows": f"{TB['conn']}-{TB['bill']},{TB['anc']}-{TB['rent']},{TB['cx']},{TB['dep']}-{TB['om']},{TB['oth']},{TB['oci']},{TB['roic']}",
                        "cols": ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"], "fmts": ["", "", "n1", "n1", "n1", "n1", "n1", "n1", "n1", "n1"],
                        "row_fmt": {**{str(TB[k]): {c: "n0" for c in "CDEFGHIJ"} for k in ("conn", "add", "avg", "lag", "bill")},
                                    **{str(TB[k]): {c: "n2" for c in "CDEFGHIJ"} for k in ("anc", "pm", "oxm")},
                                    **{str(TB[k]): {c: "p1" for c in "CDEFGHIJ"} for k in ("om", "roic")}}}},
            {"type": "note", "parts": ["@Thesis!A2"]},
            {"type": "table", "caption": f"@Thesis!A{R['RPO'][0]-2}", "grid": True,
             "head": [f"@Thesis!A{R['RPO'][0]-1}", f"@Thesis!B{R['RPO'][0]-1}", f"@Thesis!C{R['RPO'][0]-1}", f"@Thesis!D{R['RPO'][0]-1}", f"@Thesis!E{R['RPO'][0]-1}"],
             "rowgen": {"sheet": "Thesis", "rows": f"{R['RPO'][0]}-{R['RPO'][1]}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "", "n2", "n2", "n2"],
                        "row_fmt": {str(R['RPO'][0]): {"C": "n1", "D": "n1", "E": "n1"}}}},
            {"type": "kv", "items": [[f"@Thesis!A{R['CAL']+2}", f"@Thesis!C{R['CAL']+2}|p1"], [f"@Thesis!A{R['CAL']+3}", f"@Thesis!C{R['CAL']+3}|p1"],
                                     [f"@Thesis!A{R['CAL']+4}", f"@Thesis!C{R['CAL']+4}|p1"]]}]},
        {"id": "pl", "title": "三情境損益（分部 → 合併）", "blocks": [
            {"type": "pl", "years": "@PL!D3:I3", "metrics": [
                met("合併營收（十億美元）", "rev", "n1", chart=True),
                met("Cloud Infrastructure（十億美元）", "oci", "n1", chart=True),
                met("其中 OCI AI 算力（命題子列）", "ai", "n1"),
                met("分部：Cloud and software（十億美元）", "cs", "n1"),
                met("雲端與軟體分部利潤（十億美元）", "mcs", "n1"),
                met("OCI AI 營業利益率", "ociom", "p1", chart=True),
                met("營業利益率（GAAP）", "oim", "p1"),
                met("EPS（GAAP 稀釋，美元）", "eps", "n2", chart=True)]},
            {"type": "table", "caption": "營收線與分部（情境 B）", "grid": True,
             "head": ["@PL!A3", "@PL!C3", "@PL!D3", "@PL!E3", "@PL!F3", "@PL!G3", "@PL!H3", "@PL!I3"],
             "rowgen": {"sheet": "PL", "rows": f"{PB['oci']}-{PB['rev']}", "cols": ["A", "C", "D", "E", "F", "G", "H", "I"], "fmts": ["", "n2", "n2", "n2", "n2", "n2", "n2", "n2"]}},
            {"type": "table", "caption": "分部核對（三情境合計絕對差）", "grid": True,
             "head": ["@PL!A3", "@PL!C3", "@PL!D3", "@PL!E3", "@PL!F3", "@PL!G3", "@PL!H3", "@PL!I3"],
             "rowgen": {"sheet": "PL", "rows": f"{R['CHK'][0]}-{R['CHK'][2]}", "cols": ["A", "C", "D", "E", "F", "G", "H", "I"], "fmts": ["", "n3", "n3", "n3", "n3", "n3", "n3", "n3"]}},
            {"type": "h3", "text": "現金流與融資（收支模組）"},
            {"type": "table", "grid": True, "head": ["@CF!A3", "@CF!C3", "@CF!D3", "@CF!E3"],
             "rowgen": {"sheet": "CF", "rows": f"{G}-{G+6}", "cols": ["A", "C", "D", "E"], "fmts": ["", "n1", "n1", "n1"], "row_fmt": {str(G + 4): {"C": "n2", "D": "n2", "E": "n2"}}}},
            {"type": "pl", "years": "@CF!D13:I13", "metrics": [
                met("自由現金流（十億美元）", None, "n1", chart=True, rows={s: f"@CF!D{CR[s]['fcf']}:I{CR[s]['fcf']}" for s in "ABC"}),
                met("自由現金流扣股利（十億美元）", None, "n1", rows={s: f"@CF!D{CR[s]['fcfd']}:I{CR[s]['fcfd']}" for s in "ABC"}),
                met("資本支出（十億美元）", None, "n1", chart=True, rows={s: f"@CF!D{CR[s]['cx']}:I{CR[s]['cx']}" for s in "ABC"}),
                met("調整後槓桿（倍）", None, "n2", rows={s: f"@CF!D{CR[s]['lev']}:I{CR[s]['lev']}" for s in "ABC"})]},
            {"type": "note", "parts": ["@CF!A1"]}]},
        {"id": "valuation", "title": "估值", "blocks": [
            {"type": "bars", "title": "目標價與現價", "fmt": "n1", "items": [
                {"label": f"@Val!A{CON['tp']}", "value": f"@Val!C{CON['tp']}", "series": 0},
                {"label": f"@Val!A{CON['dcf']}", "value": f"@Val!C{CON['dcf']}", "series": 1},
                {"label": f"@Val!A{CON['comp']}", "value": f"@Val!C{CON['comp']}", "series": 2},
                {"label": f"@Val!A{CON['cons']}", "value": f"@Val!C{CON['cons']}", "series": 1}],
             "ref": {"label": f"@Val!A{CON['px']}", "value": f"@Val!C{CON['px']}"}},
            {"type": "table", "grid": True, "head": [f"@Val!A{SO['tps']-1}", f"@Val!C{SO['tps']-1}", f"@Val!D{SO['tps']-1}", f"@Val!E{SO['tps']-1}"],
             "rows": [[f"@Val!A{SO['tps']}", f"@Val!C{SO['tps']}:E{SO['tps']}|n2"], [f"@Val!A{DC['v12']}", f"@Val!C{DC['v12']}:E{DC['v12']}|n2"],
                      [f"@Val!A{SO['v12']}", f"@Val!C{SO['v12']}:E{SO['v12']}|n2"]]},
            {"type": "table", "caption": f"@Val!A{SO['tps']-2}", "grid": True, "head": [f"@Val!A{CON['dcf']-1}", f"@Val!C{CON['dcf']-1}", f"@Val!F{CON['dcf']-1}"],
             "rowgen": {"sheet": "Val", "rows": f"{CON['dcf']}-{CON['compm']}", "cols": ["A", "C", "F"], "fmts": ["", "n2", ""],
                        "row_fmt": {str(CON[k]): {"C": "sp1"} for k in ("up", "upc", "gcons")} | {str(CON["tvsh"]): {"C": "p1"}, str(CON["tier"]): {"C": ""}, str(CON["rec"]): {"C": ""}}}},
            {"type": "table", "caption": "@Val!A14", "grid": True, "head": ["@Val!A15", "@Val!B15", "@Val!C15", "@Val!D15", "@Val!E15"],
             "rowgen": {"sheet": "Val", "rows": f"{DC['2027']}-{DC['v12']}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "n2", "n1", "n1", "n1"],
                        "row_fmt": {str(DC["tvsh"]): {"C": "p1", "D": "p1", "E": "p1"}, str(DC["ps"]): {"C": "n2", "D": "n2", "E": "n2"}, str(DC["v12"]): {"C": "n2", "D": "n2", "E": "n2"}}}},
            {"type": "kv", "items": [["@Val!A5", "@Val!C5|p2"], ["@Val!A9", "@Val!C9|p2"], ["@Val!A10", "@Val!C10|n1"], ["@Val!A11", "@Val!C11|n3"]]},
            {"type": "table", "caption": "@Val!A30", "grid": True, "head": ["@Val!A31", "@Val!B31", "@Val!C31", "@Val!D31", "@Val!E31"],
             "rowgen": {"sheet": "Val", "rows": f"{VAL['PEERS'][0][0]}-{VAL['PEERS'][0][1]},{VAL['PEERS'][1][0]}-{VAL['PEERS'][1][1]}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "", "n1", "x1", ""]}},
            {"type": "table", "grid": True, "head": [f"@Val!A{SO['e_oci']-1}", f"@Val!C{SO['e_oci']-1}", f"@Val!D{SO['e_oci']-1}", f"@Val!E{SO['e_oci']-1}"],
             "rowgen": {"sheet": "Val", "rows": f"{SO['e_oci']}-{SO['oci_sh']}", "cols": ["A", "C", "D", "E"], "fmts": ["", "n1", "n1", "n1"],
                        "row_fmt": {str(SO["sh"]): {"C": "n3", "D": "n3", "E": "n3"}, str(SO["ps"]): {"C": "n2", "D": "n2", "E": "n2"}, str(SO["v12"]): {"C": "n2", "D": "n2", "E": "n2"},
                                    str(SO["oci_sh"]): {"C": "p1", "D": "p1", "E": "p1"}, str(SO["g_sh"]): {"C": "p1", "D": "p1", "E": "p1"}}}},
            {"type": "note", "parts": [f"@Val!A{CON['meth']}", "：", f"@Val!C{CON['meth']}"]}]},
        {"id": "recon", "title": "對帳：共識、公司指引、命題檢查", "blocks": [
            {"type": "table", "head": ["@Recon!A3", "@Recon!B3", "@Recon!C3", "@Recon!D3", "@Recon!E3", "@Recon!F3", "@Recon!G3", "@Recon!H3", "@Recon!I3"],
             "rowgen": {"sheet": "Recon", "rows": None, "cols": ["A", "B", "C", "D", "E", "F", "G", "H", "I"], "fmts": ["", "n1", "n1", "sp1", "n1", "sp1", "n1", "sp1", ""], "row_fmt": {}}},
            {"type": "note", "parts": ["@Recon!A1"]}]},
        {"id": "sens", "title": "敏感度", "blocks": [
            {"type": "tornado", "title": "加權目標價相對基準：輸入取低值", "sheet": "Sens", "rows": f"{R['SENS'][0]}-{R['SENS'][1]}", "label_col": "A", "value_col": "D", "fmt": "sp1"},
            {"type": "tornado", "title": "加權目標價相對基準：輸入取高值", "sheet": "Sens", "rows": f"{R['SENS'][0]}-{R['SENS'][1]}", "label_col": "A", "value_col": "F", "fmt": "sp1"},
            {"type": "table", "collapse": True, "summary": "展開：全部案例", "grid": True,
             "head": ["@Sens!A3", "@Sens!B3", "@Sens!C3", "@Sens!D3", "@Sens!E3", "@Sens!F3", "@Sens!H3", "@Sens!I3", "@Sens!J3", "@Sens!K3"],
             "rowgen": {"sheet": "Sens", "rows": f"{R['SENS'][0]}-{R['SENS'][1]}", "cols": ["A", "B", "C", "D", "E", "F", "H", "I", "J", "K"],
                        "fmts": ["", "", "n2", "sp1", "n2", "sp1", "p1", "p1", "n1", "n1"]}},
            {"type": "table", "caption": f"@Sens!A{R['GRID'][0]}", "grid": True, "head": [f"@Sens!A{R['GRID'][0]+1}", f"@Sens!C{R['GRID'][0]+1}", f"@Sens!D{R['GRID'][0]+1}", f"@Sens!E{R['GRID'][0]+1}"],
             "rowgen": {"sheet": "Sens", "rows": f"{R['GRID'][0]+2}-{R['GRID'][1]}", "cols": ["A", "C", "D", "E"], "fmts": ["", "n1", "n1", "n1"]}},
            {"type": "note", "parts": ["@Sens!A1"]}]},
        {"id": "status", "title": "資料狀態", "blocks": [
            {"type": "kv", "items": [["模型版本", ["$version", "（", "$date", "）"]], ["來源 xlsx", "$xlsx"], ["來源 xlsx 的雜湊指紋（取代建置 commit）", "$sha"],
                                     ["Tokenomics 快照版本", "@In!H" + str(INR["he_gb300"])]]},
            {"type": "h3", "text": "已知限制"},
            {"type": "note", "parts": ["@README!A20"]},
            {"type": "h3", "text": "「未找到」清單"},
            {"type": "scan", "needle": "未找到", "collapse": True, "summary": "展開：標為「未找到」的輸入與來源",
             "sheets": {"In": ["B"], "Q": ["A"], "Val": ["A"], "Recon": ["A"]}, "empty": "沒有標為未找到的項目。"}]},
    ],
}
# Recon 的列範圍：第 4 列到最後一列
ws = wb["Recon"]
last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value)
rg = cfg["sections"][4]["blocks"][0]["rowgen"]
rg["rows"] = f"4-{last}"
for r in range(4, last + 1):
    lab = str(ws.cell(r, 1).value or "")
    if any(k in lab for k in ("EPS", "利潤率", "ROIC", "殘差", "m_sw")):
        rg["row_fmt"][str(r)] = {"B": "n2" if "EPS" in lab else "p1", "C": "n2" if "EPS" in lab else "p1", "E": "n2" if "EPS" in lab else "p1", "G": "n2" if "EPS" in lab else "p1"}
    if "金額計" in lab or "殘差" in lab or "m_sw" in lab or "ROIC" in lab or "毛利率" in lab:
        rg["row_fmt"].setdefault(str(r), {}).update({"D": "n1" if "金額計" in lab else "sp1", "F": "n1" if "金額計" in lab else "sp1", "H": "n1" if "金額計" in lab else "sp1"})
json.dump(cfg, open(os.path.join(HERE, "html.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("html.json written", X)
