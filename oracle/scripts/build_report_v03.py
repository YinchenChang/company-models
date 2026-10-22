# Oracle v0.3 步驟 8：對照 Excel docs/reports/20261022_oracle_v0.3_TK同步.xlsx
# 數值取自 tkanchor.py、tk_sens.json、out/attrib_v03.json（node scripts/attrib_v03.js）、Tokenomics v5.24／v19.0 成品與 dist/ v0.3 成品；
# 只作報告，不進模型。用法：python3 scripts/build_report_v03.py [輸出.xlsx]
import json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, REPO)
import calendar_q, tkanchor
co = calendar_q.load(REPO); _, tk = tkanchor.load(co); P = co['periods']; K3 = ['low', 'base', 'high']
LB = co['scenarios']['labels']; AM = co['pricing']['anchorMultiple']; G = co['fleet']['generations']
TS = json.load(open(os.path.join(REPO, 'tk_sens.json'), encoding='utf-8'))
AT = json.load(open(os.path.join(REPO, 'out', 'attrib_v03.json'), encoding='utf-8'))
dist = [f for f in os.listdir(os.path.join(REPO, 'dist')) if f.endswith('.xlsx')][0]
X = load_workbook(os.path.join(REPO, 'dist', dist), data_only=True)
TKREPO = os.environ.get('TKREPO', '/home/claude/yinchenchang/tokenomics/model')
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864'); B = Font(bold=True)
wb = Workbook(); first = True


def sheet(name, head, rows, note=None, w0=52):
    global first
    ws = wb.active if first else wb.create_sheet(); first = False; ws.title = name
    r = 1
    if note: ws.cell(1, 1, note).font = Font(italic=True, size=9); r = 2
    for j, h in enumerate(head): c = ws.cell(r, 1 + j, h); c.font = H; c.fill = HF
    for row in rows:
        r += 1
        for j, v in enumerate(row): ws.cell(r, 1 + j, v)
    ws.column_dimensions['A'].width = w0
    for col in 'BCDEFGHIJKLMNOPQR': ws.column_dimensions[col].width = 15
    ws.freeze_panes = ws.cell(3 if note else 2, 2)


def xrows(sh, pred, cols):
    """dist 成品某工作表中 A 欄符合 pred 的列：回傳 [A, 指定欄...]"""
    ws = X[sh]; out = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and pred(str(a), r): out.append([a] + [ws.cell(r, c).value for c in cols])
    return out


R = {sc: tkanchor.compute(co, tk, sc, co['periodYears']) for sc in K3}
# 1 每 MW 收入 v0.2 對 v0.3
rows = []
for sc in K3:
    rows.append([f'{LB[sc]}｜v0.2（寫死 scenarios.revMW）'] + [x * 1e3 for x in co['scenarios']['revMW'][sc]])
    rows.append([f'{LB[sc]}｜v0.3（錨 × k，價格基準）'] + [x * 1e3 for x in R[sc]['rev']])
    rows.append([f'{LB[sc]}｜變動'] + [R[sc]['rev'][i] * 1e3 - co['scenarios']['revMW'][sc][i] * 1e3 for i in range(5)])
sheet('每MW收入_v0.2對v0.3', ['US$m／MW-IT·年'] + P, rows, '每 MW 年收入（100% 計費時數）；v0.2＝Tokenomics v5.24 的路徑 A 成本加成與路徑 B 市場價平均、隨容量情境同向（已停用；legacy 模式可重現）')
# 2 橋接（基準）
rows = []
for it in AT['permw']['base']:
    d = it['d']
    rows.append([it['period'], it['old'], it['recompute'], d['rounding'], d['tkAnchor'], d['costPlus'], d['marketPrice'], d['genMix'], d['kEvidence'], it['new']])
sheet('每MW收入橋接_基準', ['期間', 'v0.2 寫死值', '重算（v5.24 兩路徑）', '四捨五入', 'Tokenomics 錨更新（GB300 IF_HoldEcon 12.72 → 10.89）', '成本加成路徑退場', '市價平均路徑退場', '世代組合', 'k 證據', 'v0.3'], rows,
      'US$m／MW-IT·年；各項依序相加＝v0.3（誤差 < 0.01）。成本加成與市價平均的退場是同一個動作的兩半：舊值＝兩路徑 50/50 平均，新值＝錨 × k，其中 k 的證據只用市場價（長約、現貨）')
# 3 錨與 k 逐年
rows = []
for sc in K3:
    for key, lab, f in (('anchor', '錨（IF_HoldEcon 在役世代加權，US$m）', 1e3), ('k', '定價倍數 k（隨需 10% × k_現貨 ＋ 90% × k_長約）', 1), ('rev', '錨 × k（US$m）', 1e3), ('capRatio', '上限比（÷ IF_RevGWFleet）', 1)):
        rows.append([f'{LB[sc]}｜{lab}'] + [x * f for x in R[sc][key]])
    for g in G: rows.append([f'{LB[sc]}｜平均在役占比｜{g}'] + R[sc]['share'][g])
sheet('錨與k逐年', ['項目'] + P, rows, f"Tokenomics {co['tokenomics']['version']}（{co['tokenomics']['mergeCommit'][:7]}）；k＝隨需占比 {AM['onDemandShare']['base']:.0%} × k_現貨 {AM['spot']['base']} ＋ 其餘 × k_長約 {AM['long']['base']}")
# 4 證據
rows = [[e['label'], e['gen'], e['price'], e['unit'], e['tkName'], tk[e['tkName']]['values'][e['gen']]['基準'], e['price'] / tk[e['tkName']]['values'][e['gen']]['基準'], e['use'], e['contract'], e['term'], e['tag'], e['source'], e['date'], e['url'], e['note']] for e in AM['evidence']]
rows += [[f"合約組合｜{x['label']}", '', x['value'], x['unit'], '', '', '', '', '', '', x['tag'], x['source'], x.get('date', ''), x['url'], ''] for x in AM['contractMix']]
rows += [['找不到', '', '', '', '', '', '', '', '', '', '', t, '', '', ''] for t in AM['notFound']]
sheet('k證據表', ['證據', '世代', '價格', '單位', 'Tokenomics 名稱', 'Tokenomics 基準', '倍數 k', '用途', '合約型態', '期間', '標記', '來源', '日期', '網址', '說明'], rows)


# 5 Tokenomics v5.24 → v19.0
def tkvals(path, names):
    w = load_workbook(path, data_only=True); out = {}
    for n in names:
        dn = w.defined_names.get(n)
        if dn is None: out[n] = None; continue
        sh, ref = list(dn.destinations)[0]; c = w[sh][ref.replace('$', '')]
        out[n] = [[x.value for x in r] for r in c] if isinstance(c, tuple) else c.value
    return out


NM = ['IF_HoldEcon', 'IF_CapexIT', 'IF_RevGWFleet', 'IF_DeprLifeIT', 'IF_FacilityGW', 'IF_Util']
try:
    old = tkvals(os.path.join(TKREPO, 'archive', '20261007_Tokenomics_v5.24.xlsx'), NM)
    new = tkvals(os.path.join(TKREPO, '20261022_Tokenomics_v19.0.xlsx'), NM)
    order = []  # 15 欄＝5 世代 × (低、基準、高)；世代順序同 company.json
    rows = []
    tgen = ['Hopper H100', 'GB200 NVL72', 'GB300 NVL72', 'VR200 NVL72', '第 5 世代']
    for n in NM:
        a, b = old[n], new[n]
        for gi in range(5):
            for ci, cs in enumerate(('低', '基準', '高')):
                def pick(v):
                    if v is None: return None
                    if isinstance(v, list): return v[0][gi * 3 + ci] if len(v[0]) == 15 else None
                    return v
                x, y = pick(a), pick(b)
                if x is None and y is None: continue
                rows.append([f'{n}｜{tgen[gi]}｜{cs}', x, y, (y - x) if isinstance(x, (int, float)) and isinstance(y, (int, float)) else None])
    sheet('Tokenomics_v5.24對v19.0', ['名稱｜世代｜成本情境', 'v5.24（bdb0de7）', 'v19.0（ad33e28）', '變動'], rows,
          '直接自兩版 Tokenomics 成品的定義名稱取值（欄序＝5 世代 × 低／基準／高；第 5 世代於 Oracle 不使用）。Oracle 引用基準欄；IF_DeprLifeIT 在 v5.24 無此名稱')
except Exception as e:
    sheet('Tokenomics_v5.24對v19.0', ['說明'], [[f'無法讀取 Tokenomics 成品：{e}']])
# 6 Q1 驗證
ws = X['每MW收入_錨定']; q1 = []
on = False
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if a and str(a).startswith('F｜'): on = True; continue
    if a and str(a).startswith('G｜'): break
    if on and a and a != '項目': q1.append([a] + [ws.cell(r, c).value for c in range(2, 6)])
sheet('Q1驗證', ['項目', '單位／非 AI＝0%', '數值／非 AI＝12%', '非 AI＝20%', '非 AI＝30%'], q1,
      '只作驗證、不校準 k：Q1 實現每 MW 11.62 對模型 FY27 9.77 的差距 1.85＝(iii) 世代組合＋(ii) 價格（Oracle 自身合約 k）＋(iv) 爬坡分母＋(i) 非 AI；(iv)(i) 隨非 AI 占比 n 改變')
# 7 上限五列
cap = xrows('每MW收入_錨定', lambda a, r: ('上限檢查' in a and a.startswith('目前情境')) or a.startswith('第五列') or a.startswith('上限檢查｜第五列'), range(2, 7))
sheet('收入上限_五列', ['項目', '單位／FY27', 'FY28', 'FY29', 'FY30', 'FY31'], cap, '四層（_100／_Util／_Prod／_Life）與每 MW 收入 ÷ 客戶付費 token 營收；第五列＝Oracle 每 MW 收入 ÷ 模型商每 MW 實收營收（P×Q，只作對照，不設門檻）。四層門檻 50%，最高 37.8%')
# 8 3x3
rows = []
for m, lab, f in (('tgt', '加權目標價 US$', 1), ('eq', '五期股權募資 US$bn', 1), ('gap', '融資前缺口 US$bn', 1), ('levMax', '調整後槓桿峰值（倍）', 1), ('sh', '新股數 bn 股', 1), ('fy27Rev', 'FY27 營收 US$bn', 1)):
    for sc in K3: rows.append([f'{lab}｜{LB[sc]}'] + [TS['matrix'][sc][px][m] for px in K3])
sheet('容量x價格3x3', ['指標｜容量情境', '價格低', '價格基準', '價格高'], rows,
      f"價格低／基準／高＝k_長約 {AM['long']['low']}／{AM['long']['base']}／{AM['long']['high']}、k_現貨 {AM['spot']['low']}／{AM['spot']['base']}／{AM['spot']['high']}（皆為證據價 ÷ 目前快照成本的結果）")
# 9 敏感度
rows = [['基準（目前輸入）'] + [TS['matrix'][sc]['base'][m] for m in ('tgt', 'gap', 'eq', 'levMax', 'fy27Rev') for sc in K3]]
rows += [[x['label']] + [x['res'][sc][m] for m in ('tgt', 'gap', 'eq', 'levMax', 'fy27Rev') for sc in K3] for x in TS['sens']]
hd = ['變動'] + [f'{n}｜{LB[sc][:2]}' for n in ('目標價', '融資前缺口', '股權募資', '槓桿峰值', 'FY27 營收') for sc in K3]
sheet('敏感度', hd, rows, '價格基準；保守／基準／積極＝容量情境；US$、US$bn、倍')
# 10 目標價拆解
rows = []
for sc in K3:
    a = AT['target'][sc]; S = a['steps']
    rows.append([LB[sc], a['v02'], 0, 0, 0, a['d1a'], a['d1b'], a['d2'], a['d3'], a['d4'], a['d5'], a['d6'], a['v03'], a['total'], a['sumSteps'] - a['total'], S[0]['gap'], S[5]['gap'], S[0]['eq'], S[5]['eq'], S[0]['levMax'], S[5]['levMax']])
sheet('目標價變動拆解', ['情境', 'v0.2', '(a) 時間推移', '(b) 實際數', '(c) 假設', '(d)①a 錨取代舊推導（k＝1，仍用舊校準與舊情境價差）', '(d)①b 其餘（k＝1 的重校與舊三情境價差一併換成錨）', '(d)② 套用 k', '(d)③ 情境軸分離', '(d)④ 資本支出改引用 TK', '(d)⑤ EBITDAR 重校準', '(d)⑥ 上限檢查', 'v0.3', '總變動', '加總誤差', '融資前缺口 v0.2', '融資前缺口 v0.3', '股權募資 v0.2', '股權募資 v0.3', '槓桿峰值 v0.2', '槓桿峰值 v0.3'], rows,
      '①a＋①b＝①（快照與錨取代舊推導，k＝1）。已決定事項 12：(a)(b)(c) 由 scripts/attrib.py 確認為 0（同評價日、無新實際數、輸入不變）；(d) 由 scripts/attrib_v03.js 逐步換入（與 Excel 經 cmp31 一致），每情境各步之和＝總變動')
steps = []
for sc in K3:
    a = AT['target'][sc]
    for i, n in enumerate(('v0.2', '①快照與錨取代舊推導（k＝1）', '②套用 k', '③情境軸分離', '④資本支出改引用 TK', '⑤EBITDAR 重校準', '⑥上限檢查（＝v0.3）')):
        pass
    S = a['steps']
    for i, n in enumerate(('v0.2（起點）', '① 快照與錨取代舊推導（k＝1）', '② 套用 k', '③ 情境軸分離', '④ 資本支出改引用 TK', '⑤ EBITDAR 重校準')):
        steps.append([LB[sc], n, S[i]['tgt'], S[i]['tgt'] - S[i - 1]['tgt'] if i else None, S[i]['eq'], S[i]['gap'], S[i]['levMax'], S[i]['rev0'], S[i]['bo'], S[i]['fy27Rev']])
sheet('拆解逐步值', ['情境', '步驟（累積換入）', '目標價', '本步變動', '股權募資', '融資前缺口', '槓桿峰值', '首期每 MW 收入 US$m', '期初可計費 MW', 'FY27 營收'], steps, '各步為累積換入後的完整引擎重算；⑥上限檢查不改數字（0）')
# 11 基準 FY27–FY31 關鍵數字
inc, cf, bs, op = X['損益'], X['各期收支'], X['資產負債_新債與新股'], X['運營_產能與收入']
q1oci = co['quarterly']['actuals'] if False else None
Q1_OCI = co['scenarios']['billableRatio']['openAnnualRevenue'] / 4
rows = []
oci = [inc.cell(6, 3 + i).value for i in range(5)]; oci[0] += Q1_OCI
tot = [inc.cell(29, 3 + i).value for i in range(5)]
legacy = [tot[i] - oci[i] for i in range(5)]
ebitda = [inc.cell(34, 3 + i).value for i in range(5)]
gross = [cf.cell(6, 3 + i).value for i in range(5)]; pre = [cf.cell(8, 3 + i).value for i in range(5)]
rows += [['OCI 營收（FY27＝Q1 實際＋模型期）', *oci], ['傳統事業營收（含 Q1 實際）', *legacy], ['總營收', *tot], ['EBITDA（全年口徑）', *ebitda],
         ['毛 CapEx', *gross], ['客戶預付（抵減；FY27 自第 2 季起）', *pre], ['淨 CapEx（毛 − 預付）', *[gross[i] - pre[i] for i in range(5)]],
         ['融資前累積現金（無新融資時）', *[cf.cell(95, 3 + i).value for i in range(5)]],
         ['瀑布新債（本期舉借）', *[cf.cell(78, 3 + i).value for i in range(5)]], ['新債期末餘額', *[cf.cell(90, 3 + i).value for i in range(5)]],
         ['股權募資', *[cf.cell(84, 3 + i).value for i in range(5)]], ['新股數（bn 股）', *[bs.cell(27, 3 + i).value for i in range(5)]],
         ['租賃負債（期末）', *[bs.cell(34, 3 + i).value for i in range(5)]], ['調整後槓桿（倍；上限 4.5）', *[bs.cell(36, 3 + i).value for i in range(5)]],
         ['平均在役（計費）MW', *[op.cell(13, 3 + i).value for i in range(5)]], ['每 MW 年收入 US$m', *[R['base']['rev'][i] * 1e3 for i in range(5)]]]
sheet('基準_FY27至FY31', ['項目（US$bn，除另註）'] + P, rows, '基準情境、價格基準；自 dist/ v0.3 成品取值（Excel 重算後）。FY27 的 CapEx／預付／新債等為模型期（評價日 2026-08-31 起 0.75 年）')
# 12 指引與共識
sm = X['摘要']; rows = []
for r in range(21, 42):
    a = sm.cell(r, 1).value
    if a and str(a).startswith('差異｜'): rows.append([a, sm.cell(r, 2).value, sm.cell(r, 3).value, sm.cell(r, 4).value, sm.cell(r, 5).value])
rows.append(['FY27 營收 vs 公司指引（≥ 90）', 'US$bn', X['檢查_連動']['B8'].value, 90, X['檢查_連動']['B8'].value / 90 - 1])
sheet('指引與共識', ['項目', '單位', 'FY27', 'FY28', 'FY29'], rows, '差距 > 5% 須有原因（已決定事項 2）：見 company.json → varianceReasons；FY27 營收原因列為「已知限制」')
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, 'docs', 'reports', '20261022_oracle_v0.3_TK同步.xlsx')
wb.save(out); print('saved', out)
