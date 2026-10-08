# 以 LibreOffice headless 重算 Excel：開啟 → 全部重算 → 存檔（原格式、原路徑），再回報公式錯誤數。
# 用法：python3 scripts/recalc.py 檔案.xlsx [逾時秒數，預設 120]
# 輸出一行 JSON：{"status": "success"|"errors_found", "total_formulas", "total_errors", "error_summary"}；
# 有公式錯誤（#REF!、#DIV/0!、#VALUE!、#NAME?、#N/A）或重算失敗時以代碼 1 結束。
import sys, os, json, shutil, signal, subprocess, tempfile
from pathlib import Path
from openpyxl import load_workbook

ERRORS = ['#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A']
MACRO = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>'''
XLB = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">
<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false">
 <library:element library:name="Module1"/>
</library:library>'''
DLB = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">
<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false"/>'''


def soffice():
    for c in ('soffice', 'libreoffice'):
        if shutil.which(c): return c
    sys.exit('找不到 soffice（apt-get install -y --no-install-recommends libreoffice-calc-nogui）')


def recalc(path, timeout=120):
    path = os.path.abspath(path)
    # 每次使用獨立的暫存設定檔：不受既有 LibreOffice 程序或設定鎖定影響，巨集也不會留在使用者設定中
    prof = tempfile.mkdtemp(prefix='lo_recalc_')
    env_arg = '-env:UserInstallation=' + Path(prof).as_uri()
    # 未設 locale 時 LibreOffice 會把中文檔名轉成 ??、開不到檔而停住：固定 UTF-8，並以百分比編碼的 file URL 傳路徑
    env = dict(os.environ, LC_ALL='C.UTF-8', LANG='C.UTF-8')
    try:
        subprocess.run([soffice(), env_arg, '--headless', '--norestore', '--terminate_after_init'],
                       capture_output=True, timeout=timeout, env=env)
        std = os.path.join(prof, 'user', 'basic', 'Standard')
        os.makedirs(std, exist_ok=True)
        open(os.path.join(std, 'Module1.xba'), 'w').write(MACRO)
        open(os.path.join(std, 'script.xlb'), 'w').write(XLB)
        open(os.path.join(std, 'dialog.xlb'), 'w').write(DLB)
        mtime = os.path.getmtime(path)
        # 獨立程序群組：逾時時連同 soffice.bin 子程序一併終止，不留下卡住的程序
        pr = subprocess.Popen([soffice(), env_arg, '--headless', '--norestore',
                               'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application',
                               Path(path).as_uri()], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                              env=env, start_new_session=True)
        try:
            out, _ = pr.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            os.killpg(pr.pid, signal.SIGKILL); pr.communicate()
            raise
        r = subprocess.CompletedProcess(pr.args, pr.returncode, out, '')
        if os.path.getmtime(path) == mtime:  # 巨集未執行到存檔
            return {'status': 'recalc_failed', 'detail': (r.stdout + r.stderr).strip()[:500]}
    except subprocess.TimeoutExpired:
        return {'status': 'recalc_failed', 'detail': f'逾時 {timeout} 秒（未安裝 libreoffice-calc 時開啟 xlsx 會停住）'}
    finally:
        shutil.rmtree(prof, ignore_errors=True)

    # 公式數以公式版計，錯誤以重算後的值計
    n_formula = sum(1 for ws in load_workbook(path).worksheets for row in ws.iter_rows()
                    for c in row if isinstance(c.value, str) and c.value.startswith('='))
    summary = {}
    for ws in load_workbook(path, data_only=True).worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    e = summary.setdefault(c.value, {'count': 0, 'locations': []})
                    e['count'] += 1
                    if len(e['locations']) < 20: e['locations'].append(f'{ws.title}!{c.coordinate}')
    total = sum(e['count'] for e in summary.values())
    return {'status': 'errors_found' if total else 'success', 'total_formulas': n_formula,
            'total_errors': total, 'error_summary': summary}


if __name__ == '__main__':
    if len(sys.argv) < 2: sys.exit(__doc__ or '用法：python3 scripts/recalc.py 檔案.xlsx [逾時秒數]')
    res = recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 120)
    print(json.dumps(res, ensure_ascii=False))
    sys.exit(0 if res['status'] == 'success' else 1)
