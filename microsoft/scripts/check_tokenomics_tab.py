# MAG v0.1b（沿用 CoreWeave W1，2026-10-07）：「快照值＝Excel 分頁值」檢查（verify.sh 步驟 5d）。
# 用法：python3 scripts/check_tokenomics_tab.py 檔案.xlsx
# 檢查：(1) Excel「Tokenomics_取數」分頁每個名稱 × 世代的低成本／基準／高成本值＝company.json → tokenomics.snapshotFile 快照值（相對誤差 1e-9）；
#       (2) 每個非 missing 名稱的「基準」格都有具名範圍 TK_<名稱去掉 IF_／L1_>[_<世代代碼>]，且指向該格；
#       (3) 快照記為 missing 的名稱有列、無具名範圍；(4) W2 起其他工作表可以引用 TK_ 名稱（每 MW 改寫），只列出引用格數；
#       引用的 TK_ 名稱都必須存在（缺漏名稱應以暫代值處理，不得引用不存在的名稱）。
import json, os, re, sys
import openpyxl

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CO = json.load(open(os.path.join(REPO, 'company.json'), encoding='utf-8'))
TK = CO['tokenomics']
SNAP = json.load(open(os.path.join(REPO, TK['snapshotFile']), encoding='utf-8'))
SHEET = 'Tokenomics_取數'


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return a == b or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return (a in (None, '') and b in (None, '')) or a == b


def main(path):
    wb = openpyxl.load_workbook(path)  # 值格為常數：公式版即可讀到值
    if SHEET not in wb.sheetnames:
        print(f'失敗：{path} 沒有「{SHEET}」分頁'); return 1
    ws = wb[SHEET]
    rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    gc = {g['name']: g['code'] for g in SNAP['source']['generations']}
    short = lambda n: re.sub(r'^(IF|L1)_', '', n)
    errs, n_val, n_name = [], 0, 0
    expect_names = set()
    for nm, it in SNAP['items'].items():
        if it.get('missing'):
            k = f'TK_{short(nm)}'
            if k not in rows: errs.append(f'{nm}：missing 名稱沒有列')
            if k in wb.defined_names: errs.append(f'{k}：missing 名稱不應建具名範圍')
            continue
        if it['kind'] == 'gen_cost':
            exp = [(f'TK_{short(nm)}_{gc[g]}', [v['低成本'], v['基準'], v['高成本']], g) for g, v in it['values'].items()]
        else:
            rg = it.get('range') or {}
            exp = [(f'TK_{short(nm)}', [rg.get('低'), it['values'], rg.get('高')] if rg else [None, it['values'], None], '—')]
        for k, vals, g in exp:
            r = rows.get(k)
            if not r: errs.append(f'{k}：分頁沒有此列'); continue
            if ws.cell(r, 2).value != nm or ws.cell(r, 5).value != g:
                errs.append(f'{k}：名稱／世代欄不符（{ws.cell(r, 2).value}／{ws.cell(r, 5).value}）')
            for j, v in enumerate(vals):
                x = ws.cell(r, 6 + j).value
                n_val += 1
                if not same(v, x): errs.append(f'{k} 第 {6 + j} 欄：快照 {v!r} ≠ Excel {x!r}')
            expect_names.add(k)
            dn = wb.defined_names.get(k)
            if dn is None: errs.append(f'{k}：沒有具名範圍'); continue
            if dn.attr_text.replace("'", '') != f"{SHEET}!$G${r}": errs.append(f'{k}：具名範圍指向 {dn.attr_text}，應為 {SHEET}!G{r}')
            n_name += 1
    extra = [n for n in wb.defined_names if n.startswith('TK_') and n not in expect_names]
    if extra: errs.append(f'多出的 TK_ 具名範圍：{extra[:5]}')
    pat = re.compile(r'Tokenomics_取數|\bTK_[A-Za-z0-9_]+')
    refs = [f'{s.title}!{c.coordinate}' for s in wb.worksheets if s.title != SHEET for row in s.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith('=') and pat.search(c.value)]
    used = {m for s in wb.worksheets if s.title != SHEET for row in s.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith('=') for m in re.findall(r'\bTK_[A-Za-z0-9_]+', c.value)}
    bad = sorted(u for u in used if u not in wb.defined_names)
    if bad: errs.append(f'公式引用了不存在的 TK_ 名稱：{bad[:5]}')
    if errs:
        print(f'快照值＝Excel 分頁值：失敗 {len(errs)} 項'); [print('  ' + e) for e in errs[:20]]; return 1
    nm = sum(1 for v in SNAP['items'].values() if v.get('missing'))
    print(f'快照值＝Excel 分頁值：{len(SNAP["items"])} 個名稱（missing {nm}）、{n_val} 個值、{n_name} 個具名範圍一致；其他工作表 {len(refs)} 格公式引用 {len(used)} 個 TK_ 名稱（{SNAP["source"]["version"]}）')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
