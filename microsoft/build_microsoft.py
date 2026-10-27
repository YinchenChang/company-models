# Microsoft（MSFT）個股模型 v0.3：命題模組＋官方營收分部 → 損益表（MAG v0.3 工作單 20261027）
# 用法：python3 microsoft/build_microsoft.py raw.xlsx ；再 soffice --headless --calc --convert-to xlsx --outdir microsoft/dist raw.xlsx
#       （成品名 microsoft/dist/20261027_Microsoft個股模型_v0.3.xlsx；本程式同時寫 microsoft/html.json）
# 財年結束於 6/30（FY2027＝2026-07～2027-06）。金額：十億美元（US$B）；MW＝IT 關鍵電力 MW；股數：十億股。
# 結構：README → In → Q → Thesis → Seg → PL → CF → Val → Recon → Sens → Sources（順序固定）。
# 規格：docs/plan/個股模型原則_命題加營收分部.md；工作單：mag/docs/workorders/20261027_MAG_v0.3_個股模型重建_分部架構.md。
# 全部活公式（藍字輸入、黑字公式）；三情境 A（樂觀）／B（基準）／C（保守）與敏感度共用同一個引擎函式（engine），確保一致。
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FACTS_F = "microsoft/data/microsoft_facts_20261008.json"
SNAP_F = "microsoft/data/tokenomics_snapshot_v21.4.json"
CONS_F = "microsoft/data/consensus_msft_20261008.json"
FACTS = json.load(open(os.path.join(ROOT, FACTS_F), encoding="utf-8"))
F = {i["id"]: i for g in FACTS["groups"].values() for i in g}
SNAP = json.load(open(os.path.join(ROOT, SNAP_F), encoding="utf-8"))
GEN = {"h100": "Hopper H100", "gb200": "GB200 NVL72", "gb300": "GB300 NVL72", "vr200": "VR200 NVL72"}


def tk(name, g):
    return SNAP["items"][name]["values"][GEN[g]]["基準"]


def fb(i):
    """事實總帳 USD M → US$B"""
    return F[i]["value"] / 1000.0


BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HFILL = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); CHK = PatternFill("solid", fgColor="E2EFDA"); SUBF = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")
wb = openpyxl.Workbook()
SHEETS = ["README", "In", "Q", "Thesis", "Seg", "PL", "CF", "Val", "Recon", "Sens", "Sources"]
WS = {}
for i, n in enumerate(SHEETS):
    WS[n] = wb.active if i == 0 else wb.create_sheet(n)
    WS[n].title = n


def hdr(ws, row, labels, c0=1):
    for j, h in enumerate(labels, c0):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HFILL


def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


SC = ("A", "B", "C")
SCN = {"A": "情境 A（樂觀）", "B": "情境 B（基準）", "C": "情境 C（保守）"}
YRS = list(range(2024, 2033))
FCY = list(range(2027, 2033))


def col(y):
    return get_column_letter(3 + y - 2024)          # C＝FY2024 … K＝FY2032


YHDR = ["項目", "單位"] + [f"FY{y}" + ("A" if y <= 2026 else "") for y in YRS] + ["說明"]

# =====================================================================================
# In：輸入（一列一鍵；E＝基準、D＝低、F＝高）
# =====================================================================================
IN = WS["In"]
IN["A1"] = "In — 輸入（藍字＝輸入、黑字＝由其他輸入推得）。基準欄 E，低 D、高 F；金額十億美元（US$B），MW＝IT 關鍵電力。標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 一律給區間"
IN["A2"] = "「沿用 v0.2.2」＝MAG v0.2.2（2026-10-25，PR #105）已核准的參數與證據，原出處見來源欄；SEC 一手數字取自 microsoft/data/microsoft_facts_20261008.json（v0.1a 起的事實總帳）"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
ROWS = []
sec = lambda t: ROWS.append(("sec", t))
inp = lambda *a: ROWS.append(a)
V22 = "沿用 v0.2.2（microsoft/legacy_v0.2/company.json；交接檔 docs/handoff/20261025_Microsoft收支模型_交接檔_v0_2_2.md）"
FX = f"事實總帳 {FACTS_F}"
SA = "https://stockanalysis.com/stocks/msft/"

sec("A. 股價、股數與資本結構（評價日 2026-06-30＝FY2026 年底；目標價時點＝今日 2026-10-27 後 12 個月）")
inp("px", "現價（2026-10-26 收盤）", "US$", None, 532.60, None, "Verified", SA + "statistics/（2026-10-26 收盤 532.60，+8.09／+1.54%；擷取 2026-10-27）", "markets/invest_map/data/prices.csv：2026-10-23 收盤 524.51（stockscan；bullfincher）＋ 10-26 漲 8.09＝532.60，兩來源一致")
inp("sh_dil26", "FY2026 稀釋加權股數", "十億股", None, fb("is.sharesDilW.fy2026"), None, "Verified", FX + "｜is.sharesDilW.fy2026（10-K FY2026）", "DCF 每股（評價日 2026-06-30）的分母")
inp("sh_basic26", "FY2026 基本加權股數", "十億股", None, fb("is.sharesBasicW.fy2026"), None, "Verified", FX + "｜is.sharesBasicW.fy2026", "—")
inp("dil_gap", "稀釋與基本股數差", "十億股", None, "=sh_dil26-sh_basic26", None, "Derived", "FY2026 稀釋 − 基本", "預測期稀釋股數＝期間平均流通股數＋此差額")
inp("sh_out26", "流通股數（2026-06-30）", "十億股", None, fb("sh.out.fy2026"), None, "Verified", FX + "｜sh.out.fy2026（10-K 附註 15）", "—")
inp("cash_open", "現金與短期投資（2026-06-30）", "US$B", None, fb("bs.cashSti.fy2026"), None, "Verified", FX + "｜bs.cashSti.fy2026", "—")
inp("debt_face", "債務面額（2026-06-30）", "US$B", None, fb("debt.face.fy2026"), None, "Verified", FX + "｜debt.face.fy2026（附註 10）", "帳面值 40.294；WACC 與淨現金用面額")
for y, i in zip(range(2027, 2033), ("fy2027", "fy2028", "fy2029", "fy2030", "fy2031", None)):
    inp(f"debt_mat_{y % 100}", f"債務到期 FY{y}", "US$B", None, fb(f"debt.maturity.{i}") if i else 0.0, None, "Verified" if i else "Assumed",
        FX + (f"｜debt.maturity.{i}（附註 10 到期表）" if i else "｜到期表只列 FY2031 以後合計 34.385；FY2032 取 0"), "—")
inp("fl_liab", "融資租賃負債（2026-06-30，在帳）", "US$B", None, fb("lease.finLiab.fy2026"), None, "Verified", FX + "｜lease.finLiab.fy2026（附註 13）", "在帳融資租賃的資產已在過去的資本支出（公司口徑）內，其本金在預測期照付（CF 頁）")
inp("fl_rate", "融資租賃加權利率", "%", None, F["lease.wtdRate.fin"]["value"], None, "Verified", FX + "｜lease.wtdRate.fin", "—")
for y, i in zip(range(2027, 2033), ("fy2027", "fy2028", "fy2029", "fy2030", "fy2031", None)):
    inp(f"fl_pay_{y % 100}", f"在帳融資租賃付款 FY{y}（未折現，含息）", "US$B", None, fb(f"lease.maturity.fin.{i}") if i else "=fl_pay_31", None, "Verified" if i else "Assumed",
        FX + (f"｜lease.maturity.fin.{i}" if i else "｜FY2032 比照 FY2031（之後合計 55.490 分多年）"), "—")
inp("min_cash", "最低現金", "US$B", 30, 50, 70, "Assumed", V22 + "（defaults.minCash 50）", "低於此水準先停回購、再舉債")
inp("bb_target", "回購目標（每年）", "US$B", None, -F["cf.buyback.fy2026"]["value"] / 1000, None, "Verified", FX + "｜cf.buyback.fy2026（含員工扣稅買回 5.6）", "現金不足時先減回購")
inp("bb_g", "回購價格年增（以現價起算）", "%", 0.0, 0.08, 0.12, "Assumed", "無：約等於股東權益成本扣股利率", "只影響股數")
inp("sh_iss", "員工股票發行（每年）", "十億股", 0.035, 0.045, 0.055, "Derived", FX + "｜FY2026 回購約 0.048 十億股（計畫 0.036＋扣稅買回 5.6 ÷ 約 470 美元），基本加權股數 7.433 → 7.429 幾乎不變 → 發行約 0.045", "低／高 Assumed")
inp("dps_q0", "每季股利（FY2027 第一季付，舊額）", "US$", None, 0.91, None, "Verified", FX + "｜sh.dps.fy2026（每季 0.91）", "—")
inp("dps_q1", "每季股利（2026-09-15 宣告，新額）", "US$", None, F["sh.dpsQ.fy27q1"]["value"], None, "Interested-party", FX + "｜sh.dpsQ.fy27q1", "—")
inp("dps_27", "FY2027 每股股利（付現）", "US$", None, "=dps_q0+3*dps_q1", None, "Derived", "第一季付舊額、其後三季付新額", "—")
inp("dps_g", "每股股利年增（FY2028 起）", "%", 0.05, 0.08, 0.10, "Interested-party", V22 + "（defaults.dividend.growth；2026-09 宣告 +8%）", "—")

sec("B. 折現率與評價參數（WACC 沿用 v0.2.2 的 CAPM 輸入；權重為 [判斷] 預設，Andy 可改）")
inp("rf", "無風險利率（10 年美債，2026-10-26）", "%", None, 0.05, None, "Verified",
    "U.S. Treasury 10 年期殖利率 2026-10-26：5.00%（修正一：chat 端統一 MAG＋Oracle 個股模型的同一輸入，與 Amazon v0.3 相同）",
    FX + "｜val.ust10：2026-10-07 為 5.28%（v0.2.2 與 v0.3 初版）")
inp("beta_sa", "原始 β：StockAnalysis（5 年月資料；Yahoo 同值，推測同源，算 1 個）", "x", None, F["val.beta.sa"]["value"], None, "Verified", FX + "｜val.beta.sa（2026-10-08）", "C13：獨立來源")
inp("beta_gf", "原始 β：GuruFocus（3 年）", "x", None, F["val.beta.gurufocus"]["value"], None, "Verified", FX + "｜val.beta.gurufocus（2026-10-08）", "C13：獨立來源")
inp("blume_w", "Blume 調整：原始 β 權重（其餘向 1 收斂）", "x", None, 0.67, None, "Assumed", "共同修正工作單 F4（2026-10-27）：0.67 × 原始 β＋0.33", "—")
inp("beta", "β（兩個獨立來源 Blume 調整後的中位數）", "x", 0.95, "=MEDIAN(blume_w*beta_sa+(1-blume_w),blume_w*beta_gf+(1-blume_w))", 1.174, "Derived",
    "Blume：1.10 → 1.067、1.1741 → 1.117；中位數 1.092（修正一；v0.3 初版用原始平均 1.137）", "低 0.95＝獨立核對建議區間下緣；高＝原始 GuruFocus")
inp("erp", "股權風險溢酬", "%", 0.04, 0.05, 0.06, "Assumed", "MAG 對照表 r1 5.10（判斷，列敏感度）", V22)
inp("kd_aaa", "對照：Moody's Aaa 公司債殖利率（長天期）", "%", None, F["val.kd.aaa"]["value"], None, "Analogy", FX + "｜val.kd.aaa（FRED DAAA 2026-10-06）", "v0.3 初版 kd；Aaa 指數為 20 年以上長債，修正一改用約 10 年期")
inp("kd_spread", "約 10 年期 AAA 公司債對美債利差", "pt", 0.003, 0.006, 0.009, "Assumed", "無：MSFT 個別債券殖利率找不到；AAA 10 年期利差歷史約 0.3–0.9 個百分點", "kd＝rf＋利差（修正一 F4）")
inp("kd", "稅前債務成本（約 10 年期公司債）", "%", None, "=rf+kd_spread", None, "Derived", "rf（10 年美債）＋利差", "—")
inp("tax_w", "WACC 用稅率", "%", None, 0.20, None, "Interested-party", FX + "｜guid.fy27.etr（FY2027 指引約 20%）", "—")
inp("ke", "股東權益成本（CAPM）", "%", None, "=rf+beta*erp", None, "Derived", "rf ＋ β × ERP", "—")
inp("wacc", "WACC", "%", None, "=(px*sh_out26*ke+(debt_face+fl_liab)*kd*(1-tax_w))/(px*sh_out26+debt_face+fl_liab)", None, "Derived",
    "E／(D＋E) × ke ＋ D／(D＋E) × kd ×（1 − 稅率）；E＝現價 × 流通股數、D＝債務面額＋在帳融資租賃負債（修正一 F4）", "v0.2.2 10.90%；v0.3 初版 10.90%")
inp("g_tv", "永續成長率（FY2032 後）", "%", 0.02, 0.03, 0.04, "Assumed", "工作單 20261027 第 2 節：g 3%", "Sens ±1pt")
inp("ronic_prem", "永續期新投資報酬 − WACC", "pt", -0.02, 0.0, 0.05, "Assumed", "無：永續期新投資只賺資金成本（成長不創造價值）", "永續期自由現金流＝NOPAT ×（1 − g ÷ 新投資報酬）；替代口徑（FY2032 自由現金流 ×（1＋g））見 Sens 判斷預設表")
inp("base_fy", "評價日所屬會計年度（最新已申報年度）", "年", None, 2026, None, "Verified", "FY2026 10-K（期末 2026-06-30）", "公式中的年數以「年度列 − base_fy」計算（修正一 F6：不寫死）")
inp("mid_yr", "期中折現：年中偏移", "年", None, 0.5, None, "Assumed", "現金流平均發生在年中", "—")
inp("roll_d1", "評價日（2026-06-30）", "日期", None, "=DATE(2026,6,30)", None, "Verified", "FY2026 年底", "—")
inp("roll_d2", "目標價時點（今日 2026-10-27 後 12 個月）", "日期", None, "=DATE(2027,10,27)", None, "Assumed", "工作單：目標價約 12 個月後", "—")
inp("days_yr", "每年天數", "天", None, 365, None, "Assumed", "—", "—")
inp("t_roll", "評價日 → 目標價時點（年）", "年", None, "=(roll_d2-roll_d1)/days_yr", None, "Derived", "2026-06-30 → 2027-10-27", "DCF 每股以 WACC 推到目標價時點（已決定事項 12 沿用），並扣除期間股利（修正一 F6）")
inp("roll_div_q", "評價日到目標價時點之間、FY2028 已付的股利季數", "季", None, 1, None, "Derived", "FY2028 第一季股利約 2027-09 支付，早於 2027-10-27", "—")
inp("q_per_yr", "每年季數", "季", None, 4, None, "Verified", "—", "—")
inp("fy28_end", "FY2028 年底", "日期", None, "=DATE(2028,6,30)", None, "Verified", "財年 6 月結束", "—")
inp("w_ntm28", "目標價時點的未來 12 個月中 FY2028 的比重", "x", None, "=(fy28_end-roll_d2)/days_yr", None, "Derived", "2027-10-27 起 12 個月：FY2028 約 0.68、FY2029 約 0.32", "SOTP 用 NTM EBITDA")
inp("w_dcf", "目標價中 DCF 的權重（其餘為分部加總 SOTP）", "%", 0.3, 0.5, 0.7, "Assumed", "[判斷] 預設：工作單 20261027 第 2 節（與 AI半導體 DCF＋Comparable 各半一致；Comparable 在分部層做）", "Andy 可改")
for s, w in zip(SC, (0.25, 0.5, 0.25)):
    inp(f"w_{s}", f"情境權重 {s}", "%", None, w, None, "Assumed", "[判斷] 預設：工作單 20261027 第 2 節 A 25%／B 50%／C 25%", "Andy 可改")
inp("oai_stake", "OpenAI 持股比例", "%", None, F["oai.stake.fy2026"]["value"], 0.27, "Verified", FX + "｜oai.stake.fy2026（10-K FY2026 附註 1：approximate 25%，經查核，最新）",
    "oai.stake.recap：重組時（2025-10-28）約 27%（10-Q FY26Q1）；10-K 稱其後因其他籌資稀釋降到約 25%。修正一維持最新查核值 25%，27% 列判斷替代")
inp("oai_val", "OpenAI 最近一輪投後估值", "US$B", None, F["oai.valuation"]["value"] / 1000, None, "Interested-party", FX + "｜oai.valuation（2026-03 結案，募資 122）", "新創自報估值")
inp("liq_disc", "未上市持股流動性折價", "%", 0.0, 0.2, 0.4, "Assumed", "工作單 20261027：20%", V22)
inp("oai_v", "OpenAI 持股價值（扣折價）", "US$B", None, "=oai_stake*oai_val*(1-liq_disc)", None, "Derived", "比例 × 估值 ×（1 − 折價）", "Anthropic 持股比例未揭露 → 不估值（資料缺口）")
inp("m_ai", "Azure AI 倍數（EV／NTM EBITDA）", "x", 4, 6, 10, "Assumed", V22 + "（valuation.evEbitda 6×）", "AI 基礎設施租賃同業（CRWV TTM 25×、NBIS 不具意義、ORCL NTM 9.5×）口徑不一，沿用 v0.2.2")
inp("m_cloud", "非 AI 雲端倍數（Azure 其他）", "x", None, 13.82, None, "Derived", "microsoft/data/peers_cloud_20261008.json：ORCL、IBM、SAP、CRM、NOW、AKAM、EQIX、DOCN 中位數（C12）", "2026-10-07 收盤")
inp("m_sw", "企業軟體倍數（M365 cloud、授權、產業、Frontier）", "x", None, F["val.peer.sw.median"]["value"], None, "Derived", FX + "｜val.peer.sw.*：ORCL 9.54、SAP 15.02、CRM 10.80、ADBE 6.60、NOW 20.65 中位數（不含 MSFT）", "2026-10-07")
inp("m_ads", "數位廣告倍數（Search and advertising）", "x", None, 9.11, None, "Derived", "microsoft/data/peers_ads_20261009.json：META、PINS、TTD、APP、RDDT、SNAP 中位數（C18）", "2026-10-07")
inp("m_game", "遊戲倍數（XBOX）", "x", None, F["val.peer.game.median"]["value"], None, "Derived", FX + "｜val.peer.game.*：TTWO 50.9、NTDOY 12.63、NTES 8.23 中位數（TTM；NTM 找不到）", "2026-10-08")
inp("m_pc", "PC／裝置倍數（Windows OEM and devices）", "x", None, F["val.peer.pc.median"]["value"], None, "Derived", FX + "｜val.peer.pc.*：HPQ 7.6、DELL 22.1、Lenovo 14.47 中位數（TTM）", "2026-10-08")
inp("m_own", "Microsoft 自身 NTM EV／EBITDA（對照）", "x", None, F["val.msft.ownNtm"]["value"], None, "Derived", FX + "｜val.msft.ownNtm（2026-10-07）", "只用於判斷預設替代（非 AI 各線改用自身倍數）")
for k, lab, v, note in (("al_m365", "Microsoft 365 cloud", 0.60, "舊 PBP 分部 FY2026 59.9% 類比"), ("al_lic", "Productivity and server licensing", 0.75, "地端授權毛利高、銷售費用低"),
                        ("al_ind", "Industry solutions", 0.30, "含 Dynamics、LinkedIn 等"), ("al_fro", "Frontier and support services", 0.30, "含支援服務"),
                        ("al_srch", "Search and advertising", 0.30, "—")):
    inp(k, f"產品線營業利益率分配：{lab}（SOTP 用）", "%", v - 0.1, v, v + 0.1, "Assumed", FX + "｜alloc.oiMargin（公司不揭露產品線利潤）；" + V22, note + "；Azure 其他、XBOX＋Windows 為分部殘差")

sec("C. Tokenomics v21.4 快照（microsoft/data/tokenomics_snapshot_v21.4.json；基準成本欄；每 GW＝IT 關鍵電力；沿用 v0.2.2，本版不升級）")
TKN = "Tokenomics v21.4（20261024_Tokenomics_v21.4.xlsx，master 3ac70ac）"
for g in ("h100", "gb200", "gb300", "vr200"):
    inp(f"tk_hold_{g}", f"IF_HoldEcon｜{GEN[g]}（每 MW 年持有成本，WACC 10% 打平）", "US$M/MW-年", None, tk("IF_HoldEcon", g), None, "Derived", TKN, "基準成本欄")
for g in ("h100", "gb200", "gb300", "vr200"):
    inp(f"tk_cxit_{g}", f"IF_CapexIT｜{GEN[g]}（每 MW IT 資本支出）", "US$M/MW", None, tk("IF_CapexIT", g), None, "Derived", TKN, "—")
for g in ("h100", "gb200", "gb300", "vr200"):
    inp(f"tk_opex_{g}", f"IF_OpexGW｜{GEN[g]}（每 MW 年營運成本）", "US$M/MW-年", None, tk("IF_OpexGW", g), None, "Derived", TKN, "電力、維護、人員、稅保")
inp("tk_cxfac", "IF_CapexFacility（每 MW 機房資本支出）", "US$M/MW", None, tk("IF_CapexFacility", "vr200"), None, "Derived", TKN, "各世代相同")
inp("tk_gpuhr_h100", "IF_GPUhrEcon｜Hopper H100（每 GPU 小時打平價）", "US$/GPU-hr", None, tk("IF_GPUhrEcon", "h100"), None, "Derived", TKN, "k_現貨的分母")
inp("life_it", "AI IT 折舊年限（IF_DeprLifeIT）", "年", 5, 6, 7, "Derived", TKN + "；10-K 伺服器與網路設備 2–6 年（pl.usefulLife.servers）", "Sens ±1 年")
inp("life_fac", "機房折舊年限", "年", None, round(tk("IF_CapexFacility", "vr200") / tk("IF_DeprFac", "vr200")), None, "Derived", TKN + "（IF_CapexFacility ÷ IF_DeprFac＝20）", "公司 FY2027 起建物 15 → 25 年（guid.usefulLife）")
inp("fac_wacc", "機房租金年金折現率（Tokenomics 經濟 WACC）", "%", None, 0.10, None, "Derived", V22 + "（tokenomics.holdEconWacc：Tokenomics Inputs!E28 基準 10%）", "與 IF_HoldEcon 同一折現率")
inp("fac_rent", "租用機房每 MW 年租金（資本回收年金）", "US$M/MW-年", None, "=tk_cxfac*fac_wacc/(1-(1+fac_wacc)^(-life_fac))", None, "Derived", "IF_CapexFacility × 年金因子（fac_wacc、機房年限）", "C22：租用機房的租金計入 AI 營運成本")

sec("D. 命題模組：AI MW、對外比例、k（沿用 v0.2.2；不以 Azure 營收、RPO、AI run-rate 反推）")
for y, v, t, n in ((2024, 640, "Derived", "v0.2.2 defaults.mwYearEnd"), (2025, 1630, "Derived", "v0.2.2 defaults.mwYearEnd"), (2026, 3250, "Derived", "三法交叉：自有在役 AI 4,500 × 對外 60%＋neocloud 租用 1,100 × 50%（split.extMW）")):
    inp(f"mw_ext_{y % 100}", f"在役對外 AI MW（FY{y} 年底，含租用）", "MW", 2150 if y == 2026 else None, v, 4875 if y == 2026 else None, t, V22 + "｜" + n, "公司未揭露 AI 在役 MW（不存在）")
for y, v, t in ((2024, 150, "Assumed"), (2025, 400, "Assumed"), (2026, 550, "Derived")):
    inp(f"rent_ext_{y % 100}", f"其中租用 neocloud 的對外 AI MW（FY{y} 年底）", "MW", None, v, None, t, V22 + "（capexModel.rentedExt：期初 550＝neocloud 在役 1,100 × 對外 50%）" if y == 2026 else "無：CoreWeave 2024、2025 營收約 62%–67% 來自 Microsoft（nc.crwv.*）推估", "租用 MW 收入照算，不計資本支出與折舊（C20）")
inp("rent_ext_f", "租用 neocloud 的對外 AI MW（FY2027 起）", "MW", None, 756, None, "Derived", V22 + "（rentedExt.path：加 Nscale 對外 205.5；到期以相同條件續約）", "—")
inp("pace", "在役對外 AI MW 每年新增（三情境相同）", "MW／年", None, 2527.2, None, "Derived", V22 + "（C10：由 FY2027 資本支出 204.2 解出，scripts/calib_pace.js）", "工作單：情境差異集中在命題的變現驅動；替代（保守＝Epoch 可辨識站點 650／年）見 Sens 判斷預設表")
inp("pace_y0", "MW 新增速度開始收斂的年度（此年後逐年線性收斂）", "年", None, 2029, None, "Assumed", "共同修正工作單 F5：FY2029 之後逐年收斂，終值不得落在建置高峰", "—")
inp("pace_y1", "MW 新增速度收斂完成的年度", "年", None, 2032, None, "Assumed", "同上：FY2032 只剩與永續成長 g 一致的增量（＝g_tv × FY2029 年底在役對外 MW）", "汰換另計")
inp("k_ramp", "情境 k 由基準 k 移到情境值的年數", "年", None, 3, None, "Assumed", "無：合約重新定價需時", "—")
inp("pace_epoch", "替代：Epoch 可辨識站點速度", "MW／年", None, 650, None, "Derived", V22 + "（scenarios.mwPath.pace.low；Fairwater Wisconsin 369 → 2,263）", "上限 cap_epoch")
inp("cap_epoch", "替代：Epoch 可辨識站點上限", "MW", None, 4386, None, "Derived", V22, "—")
inp("ext_sh", "對外比例（對外 AI MW ÷ 自有 AI MW）", "%", 0.5, 0.6, 0.75, "Assumed", V22 + "（C23：公司無揭露，最大不確定）", "Sens 0.5／0.75")
inp("self_build", "機房自建比例（自建＋融資租賃）", "%", 0.4, 0.6, 0.8, "Assumed", V22 + "（capexModel.selfBuild）", "其餘為營業租賃：不計資本支出，計租金（C22）")
for k, lab, v in (("mix_h", "Hopper", 0.2), ("mix_g2", "GB200", 0.45), ("mix_g3", "GB300（含 Maia 200）", 0.35)):
    inp(k, f"期初在役世代占比：{lab}", "%", None, v, None, "Assumed", V22 + "（mw.genMix.open；區間 Hopper 15–30%、GB200 35–55%、GB300 20–40%）", "FY2026 以前各年同此占比")
for y, v in ((27, 0.7), (28, 0.3), (29, 0.1)):
    inp(f"a_g3_{y}", f"新增 MW 中 GB300（含 Maia 200）占比 FY20{y}" + ("（FY2030 起同）" if y == 29 else ""), "%", None, v, None, "Assumed", V22 + "（chips.mixAdds：GB300 60%／20%／0%＋Maia 10%；其餘 VR200）", "Maia 200 以 GB300 × 1.0 計（C2；NonNV 只登錄不切換）")
inp("cxit_open", "期初世代加權 IT 資本支出", "US$M/MW", None, "=mix_h*tk_cxit_h100+mix_g2*tk_cxit_gb200+mix_g3*tk_cxit_gb300", None, "Derived", "期初占比 × IF_CapexIT", "FY2026 以前各年 AI 資本支出估計與折舊用")
inp("p_long", "長約證據價格：IREN–Microsoft GB300 五年約", "US$M/MW-年", None, 9.7, None, "Analogy", V22 + "（kevidence lc.iren：TCV 9,700 ÷ 5 年 ÷ 200 MW）", "Microsoft 付給 neocloud 的買價，代理長約賣價")
inp("p_spot", "現貨證據價格：Silicon Data H100 指數", "US$/GPU-hr", None, 2.82, 12.29, "Verified", V22 + "（kevidence spot.h100.index，2026-10-06）", "高＝Azure H100 隨需牌價 12.29（只作上限）")
inp("k_long", "k_長約", "x", None, "=p_long/tk_hold_gb300", None, "Derived", "證據價格 ÷ 同世代 IF_HoldEcon（kspec 規格）", "v0.2.2 0.891")
inp("p_long_lo", "長約低端價格（v0.2.2 判斷倍數 0.45 × 12.762 換算）", "US$M/MW-年", None, 5.7429, None, "Assumed", V22 + "（kevidence long.low：非觀察價格）", "只用於判斷預設替代")
inp("k_long_lo", "k_長約低端", "x", None, "=p_long_lo/tk_hold_vr200", None, "Derived", "長約低端價格 ÷ IF_HoldEcon｜VR200", "v0.2.2 0.470")
inp("k_spot", "k_現貨", "x", None, "=p_spot/tk_gpuhr_h100", None, "Derived", "證據價格 ÷ IF_GPUhrEcon｜H100", "v0.2.2 1.828")
for s, v, n in (("A", 0.45, "長約 45%＋現貨 55%（v0.1a split.k 高端口徑）"), ("B", 0.70, "沿用 v0.2.2 longShare 70%"), ("C", 1.00, "全長約（v0.1a split.k 低端口徑）")):
    inp(f"ls_{s}", f"長約占比（情境 {s}）", "%", None, v, None, "Assumed", n, "[判斷] 預設：A／C 的價格軸差異是本版主要情境軸")
for s in SC:
    inp(f"k_{s}", f"k（情境 {s}）", "x", None, f"=ls_{s}*k_long+(1-ls_{s})*k_spot", None, "Derived", "長約占比 × k_長約 ＋（1 − 長約占比）× k_現貨", "B＝v0.2.2 混合 k 1.172")
inp("chip_f", "自研晶片係數（Maia 以 GB300 計）", "x", 0.7, 1.0, 1.0, "Assumed", V22 + "（C2；NonNV 持有成本比 0.4–0.6 只登錄）", "—")
inp("rent_mw", "租用 neocloud 每 MW 年租金（全額）", "US$M/MW-年", None, 8.89, None, "Derived", V22 + "（已揭露合約全額年租金 13.96 ÷ 1,571 MW）", "租用 MW 不另計營運成本")
inp("ai_net_open", "期初 AI PP&E 淨額 ÷ 毛額", "%", 0.65, 0.8, 0.9, "Assumed", V22 + "（capexModel.aiNetShare）", "AI 增量 ROIC 的期初投入資本")
inp("capex_floor27", "FY2027 資本支出下限（公司口徑，含融資租賃）", "US$B", None, F["capex.fy27.derived"]["value"] / 1000, None, "Derived", FX + "｜capex.fy27.derived：曆年 2026 約 175 −（FY26 下半年 31.9＋41.0）＝FY27 上半年 102.1，下半年＝上半年", "FY27Q1 指引 >50 吻合；共識 192.8（推定現金口徑）")

sec("E. Microsoft 365 cloud 的 Copilot（命題子列：付費座席 × 每座席月費）")
for y, v, t, n in ((2024, 1, "Assumed", "無：Copilot 2023-11 上市，FY2024 年底座席未揭露"), (2025, 5, "Assumed", "無：FY2025 年底座席未揭露（FY26Q2 15 百萬）"),
                   (2026, F["call.copilotSeats.fy26q4"]["value"], "Interested-party", FX + "｜call.copilotSeats.fy26q4（FY26Q4 法說 30 百萬席；FY26Q2 15 百萬）")):
    inp(f"cop_s{y % 100}", f"M365 Copilot 付費座席（FY{y} 年底）", "百萬席", None, v, None, t, n, "—")
inp("cop_px", "Copilot 每座席月費（實收）", "US$／月", 18, 25, 30, "Assumed", "無：牌價 30 美元／月，大客戶折扣未揭露", "Sens 18／30")
for s, v in (("A", 30), ("B", 20), ("C", 10)):
    inp(f"cop_add_{s}", f"Copilot 每年淨增座席 FY2027–FY2029（情境 {s}）", "百萬席", None, v, None, "Assumed", "FY2026 淨增約 25 百萬（5 → 30，Assumed 起點）", "FY2030 起 × cop_tail")
inp("cop_tail", "Copilot FY2030 起淨增相對前期的比例", "x", None, 0.5, None, "Assumed", "無：滲透率上升後放緩", "—")
inp("internal_ai_reading", "自用 AI 讀法（2＝主值：自用 AI MW 以 k＝1 計影子效益；1＝敏感度：自用 AI 只計成本）", "代碼", 1, 2, None, "Assumed",
    "共同修正工作單 F1（chat 端預設：讀法 2 為主值、讀法 1 為敏感度；Andy 待決）", "影子效益＝MAX(0, 自用在役 MW × 每 MW 持有成本（k＝1）− Copilot × 核心利益率)：扣 Copilot 已計的部分，避免重複")

sec("F. 產品線成長路徑：近一年年增（FY2026 對 FY2025，8-K 重編口徑）線性收斂到長期值（FY2032 到達）")
inp("n_conv", "收斂年數", "年", None, 6, None, "Assumed", "與 AI半導體 v2 DCF 明確期（到 FY2032）一致", "—")
for k, lab, lo, b, hi, n in (("glt_azo", "Azure 其他（非 AI 雲端）", 0.05, 0.08, 0.12, V22 + "（azureNonAi gLT 8%）"), ("glt_m365b", "M365 cloud 基礎（不含 Copilot）", 0.04, 0.07, 0.10, "v0.2.2 M365 含 Copilot 10%；扣 Copilot 後取 7%"),
                             ("glt_lic", "Productivity and server licensing", 0.0, 0.02, 0.04, V22), ("glt_ind", "Industry solutions", 0.04, 0.07, 0.10, V22 + "（indFrontier 7%）"),
                             ("glt_fro", "Frontier and support services", 0.04, 0.07, 0.10, V22 + "（indFrontier 7%）"), ("glt_srch", "Search and advertising", 0.04, 0.07, 0.10, V22),
                             ("glt_xbox", "XBOX", -0.02, 0.02, 0.05, V22 + "（devGaming 2%）"), ("glt_win", "Windows OEM and devices", -0.02, 0.02, 0.04, V22 + "（devGaming 2%）")):
    inp(k, f"長期年增：{lab}", "%", lo, b, hi, "Assumed", n, "三情境相同（非命題業務）；Sens 用低／高")

sec("G. 損益、現金與非 AI 資本支出")
inp("dm_seg", "分部核心利益率調整（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "加在 Agents and Infra 核心利益率與 Devices and Consumer 利益率")
inp("oie_f", "業外收支（不含持股評價損益；每年）", "US$B", -1.0, -0.4, 0.4, "Interested-party", FX + "｜guid.fy27q1.oie：FY27Q1 約 −0.1（每季）", "持股評價損益（OpenAI 權益法、Anthropic）預測期設 0")
inp("tax_f", "有效稅率（FY2027 起）", "%", 0.19, 0.20, 0.21, "Interested-party", FX + "｜guid.fy27.etr（約 20%）", "FY2026 19.4%")
inp("sbc_r", "股票薪酬 ÷ 營收", "%", None, f"=Q!M{{SBC}}/Q!M{{REV}}", None, "Derived", "FY2026 實績（Q 頁）", "—")
inp("wc27", "FY2027 營運資金變動", "US$B", None, F["cf.wc.total.fy2026"]["value"] / 1000, None, "Derived", FX + "｜cf.wc.total.fy2026（C14：首期＝上一年度實績）", "FY2028 起見 wc_f")
inp("wc_f", "FY2028 起營運資金變動（每年）", "US$B", None, 0, None, "Assumed", V22 + "（wcPctOfRevGrowth 0）", "—")
inp("life_nonai", "非 AI 資產折舊年限", "年", 8, 10, 15, "Assumed", V22 + "（C17）", "非 AI 折舊＝前期 ×（1 − 1／年限）＋ 當期非 AI 資本支出 ÷ 年限")
for k, lab, v in (("cx_azo", "Azure 其他", 0.20), ("cx_m365", "Microsoft 365 cloud", 0.10), ("cx_lic", "授權", 0.02), ("cx_ind", "Industry＋Frontier", 0.05), ("cx_srch", "Search and advertising", 0.08), ("cx_dev", "XBOX＋Windows", 0.03)):
    inp(k, f"非 AI 資本支出 ÷ 營收：{lab}", "%", None, v, None, "Assumed", V22 + "（legacyBiz.lines.cx）", "公司不揭露分部資本支出")

sec("H. 市場共識與公司指引（只作 Recon 對照，不作校準）")
CS = json.load(open(os.path.join(ROOT, CONS_F), encoding="utf-8"))
AE = CS["annualEstimates"]
inp("cons_tp", "市場共識目標價（S&P，57 位）", "US$", 440, 603.43, 870, "Interested-party", SA + "forecast/（2026-10-26 頁面；中位數 600；Strong Buy 41／Buy 13／Hold 3）", "2026-10-08 為 587.63（56 位）")
inp("cons_rev27", "共識營收 FY2027（S&P）", "US$B", 380.0, 391.6, 402.7, "Interested-party", SA + "forecast/（2026-10-26，52 位）", "MarketScreener 391.271（2026-10-08）")
inp("cons_eps27", "共識 EPS FY2027（S&P，non-GAAP）", "US$", 19.06, 19.76, 20.70, "Interested-party", SA + "forecast/（2026-10-26）", "MarketScreener GAAP 19.72")
for y in (2028, 2029):
    d = AE[f"FY{y % 100}"]
    inp(f"cons_rev{y % 100}", f"共識營收 FY{y}（MarketScreener）", "US$B", None, d["revenue"], None, "Interested-party", CONS_F + "（2026-10-08）", "—")
for y in (2027, 2028, 2029):
    d = AE[f"FY{y % 100}"]
    inp(f"cons_ebit{y % 100}", f"共識營業利益 FY{y}（MarketScreener）", "US$B", None, d["ebit"], None, "Interested-party", CONS_F + "（2026-10-08）", "GAAP 營業利益口徑")
    inp(f"cons_cx{y % 100}", f"共識資本支出 FY{y}（MarketScreener，推定現金口徑）", "US$B", None, d["capex"], None, "Interested-party", CONS_F, "模型為公司口徑（含融資租賃）")
for y in (2028, 2029):
    d = AE[f"FY{y % 100}"]
    inp(f"cons_eps{y % 100}", f"共識 EPS FY{y}（MarketScreener，GAAP）", "US$", None, d["eps"], None, "Interested-party", CONS_F, "—")
inp("g_q1_rev", "FY27Q1 營收指引", "US$B", 89.85, 90.4, 90.95, "Interested-party", FX + "｜guid.fy27q1.rev（8-K 2026-09-02）", "—")
inp("g_q1_ai", "FY27Q1 Agents and Infra 營收指引", "US$B", 75.15, 75.45, 75.75, "Interested-party", FX + "｜guid.fy27q1.rev.ai", "—")
inp("g_q1_dc", "FY27Q1 Devices and Consumer 營收指引", "US$B", 14.7, 14.95, 15.2, "Interested-party", FX + "｜guid.fy27q1.rev.dc", "—")
inp("g_q1_az", "FY27Q1 Azure 年增指引", "%", 0.44, 0.445, 0.45, "Interested-party", FX + "｜guid.fy27q1.azure", "—")
inp("g_q1_cx", "FY27Q1 資本支出指引（下限）", "US$B", None, 50, None, "Interested-party", FX + "｜guid.fy27q1.capex（over 50）", "—")
inp("g_fy27_opm", "FY2027 營業利益率變動指引（上限）", "pt", None, -0.01, None, "Interested-party", FX + "｜guid.fy27.opm（down less than a point）", "—")
inp("ai_runrate", "公司 AI 業務年化營收（FY26Q3）", "US$B", None, 37, None, "Interested-party", FX + "｜call.aiRunRate.fy26q3（含 Copilot 等自有 AI 產品）", "C3：只作驗證列")

sec("I. v0.2.2 對照（microsoft/dist/_待刪除/20261025_Microsoft收支模型_v0_2_2.xlsx 評價_DCF與目標價）")
inp("old_tp", "v0.2.2 基準目標價（讀法 1）", "US$", 360.16, 359.06, 380.56, "Derived", "v0.2.2 成品：低欄＝積極、高欄＝保守情境", "讀法 2：429.43")
inp("old_dcf", "v0.2.2 DCF 每股（推到目標價時點）", "US$", None, 324.13, None, "Derived", "v0.2.2 成品", "—")
inp("old_ev", "v0.2.2 EV／EBITDA 每股（融資後）", "US$", None, 387.64, None, "Derived", "v0.2.2 成品", "—")
inp("old_w", "v0.2.2 DCF 權重", "%", None, 0.45, None, "Derived", "v0.2.2（DCF 45%＋分部 EV/EBITDA 55%）", "—")

IR = {}
r = 4
for srow in ROWS:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")):
                IN.cell(r, j).font = BLUE
            if unit in ("%", "pt"):
                IN.cell(r, j).number_format = "0.0%"
            elif unit == "x":
                IN.cell(r, j).number_format = "0.000"
            elif unit == "日期":
                IN.cell(r, j).number_format = "yyyy-mm-dd"
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}")
    if lo is not None:
        defname(key + "_lo", f"In!$D${r}")
    if hi is not None:
        defname(key + "_hi", f"In!$F${r}")
    if key.startswith(("w_", "w_dcf")):
        for c in range(1, 10): IN.cell(r, c).fill = KEY
    IR[key] = r; r += 1
for c_, w in zip("ABCDEFGHI", [14, 60, 12, 10, 12, 10, 16, 80, 70]):
    IN.column_dimensions[c_].width = w
IN.freeze_panes = "D4"

# =====================================================================================
# Q：實績（FY25Q1–FY26Q4 八季＋FY2025、FY2026 兩個完整年度；8-K 2026-09-02 重編口徑）
# =====================================================================================
Q = WS["Q"]
Q["A1"] = "Q — 實績（十億美元）。分部與產品線為 FY2027 起的報導口徑（8-K 2026-09-02「Segment History／Products and Service Offerings — As Restated」，未經查核）；合併損益、現金流量取自 10-K／財報新聞稿（事實總帳）"
Q["A2"] = "J、K 欄＝四季合計（公式）；L、M 欄＝公司揭露的年度值（輸入）；核對列實績必須為 0。舊三分部（PBP／IC／MPC）列在下方對照（FY2026 以前口徑）"
QY = ["fy25q1", "fy25q2", "fy25q3", "fy25q4", "fy26q1", "fy26q2", "fy26q3", "fy26q4"]
QC = "BCDEFGHI"
hdr(Q, 3, ["項目", "FY25Q1（～2024-09）", "FY25Q2（～12）", "FY25Q3（～2025-03）", "FY25Q4（～06）", "FY26Q1（～2025-09）", "FY26Q2（～12）", "FY26Q3（～2026-03）", "FY26Q4（～06）",
           "FY2025（四季合計）", "FY2026（四季合計）", "FY2025（公司年度值）", "FY2026（公司年度值）", "標記", "來源（事實總帳 id）"])
QR = {}
qr = [4]


def qrow(key, label, qids=None, aids=None, qform=None, sumcols=True, fmt="#,##0.000", tag="Verified", src="", neg=False):
    r_ = qr[0]; Q.cell(r_, 1, label)
    sg = -1 if neg else 1
    for j, c in enumerate(QC):
        if qids:
            v = F[qids.format(QY[j])]["value"] * sg / 1000.0
            Q[f"{c}{r_}"] = v; Q[f"{c}{r_}"].font = BLUE
        elif qform:
            Q[f"{c}{r_}"] = "=" + qform(c)
        Q[f"{c}{r_}"].number_format = fmt
    if sumcols and (qids or qform):
        Q[f"J{r_}"] = f"=SUM(B{r_}:E{r_})"; Q[f"K{r_}"] = f"=SUM(F{r_}:I{r_})"
    elif qform and not sumcols:
        pass
    if aids:
        for c, a in zip("LM", aids):
            if a is None:
                continue
            if isinstance(a, str) and a.startswith("="):
                Q[f"{c}{r_}"] = a
            else:
                Q[f"{c}{r_}"] = F[a]["value"] * sg / 1000.0; Q[f"{c}{r_}"].font = BLUE
    for c in "JKLM":
        Q[f"{c}{r_}"].number_format = fmt
    Q.cell(r_, 14, tag); Q.cell(r_, 15, src)
    QR[key] = r_; qr[0] += 1
    return r_


def qsec(t):
    Q.cell(qr[0], 1, t).font = BOLD; qr[0] += 1


def qcheck(key, label, f):
    r_ = qr[0]; Q.cell(r_, 1, label)
    for c in QC + "JKLM":
        e = f(c)
        if e is not None:
            Q[f"{c}{r_}"] = "=" + e; Q[f"{c}{r_}"].number_format = "0.000"
    for c in range(1, 14): Q.cell(r_, c).fill = CHK
    Q.cell(r_, 14, "核對"); QR[key] = r_; qr[0] += 1


R8 = "8-K 2026-09-02 附件 99.1（As Restated）"
qsec("1. 官方分部與產品線營收（FY2027 起報導口徑）")
qrow("ai", "Agents and Infra", "q.seg27.ai.rev.{}", ("seg27.ai.rev.fy2025", "seg27.ai.rev.fy2026"), src=R8 + "｜q.seg27.ai.rev.*、seg27.ai.rev.*")
for k, lab in (("azure", "  Azure"), ("m365", "  Microsoft 365 cloud"), ("lic", "  Productivity and server licensing"), ("ind", "  Industry solutions"), ("fro", "  Frontier and support services")):
    pid = {"azure": "azure", "m365": "m365cloud", "lic": "licensing", "ind": "industry", "fro": "frontier"}[k]
    qrow(k, lab, f"q.prod27.{pid}.{{}}", (f"prod27.{pid}.fy2025", f"prod27.{pid}.fy2026"), src=R8 + f"｜q.prod27.{pid}.*")
qcheck("chk_ai", "檢查：Agents and Infra 產品線合計 − 分部", lambda c: f"SUM({c}{QR['azure']}:{c}{QR['fro']})-{c}{QR['ai']}")
qrow("dc", "Devices and Consumer", "q.seg27.dc.rev.{}", ("seg27.dc.rev.fy2025", "seg27.dc.rev.fy2026"), src=R8 + "｜q.seg27.dc.rev.*")
for k, lab in (("srch", "  Search and advertising"), ("xbox", "  XBOX"), ("win", "  Windows OEM and devices")):
    pid = {"srch": "searchads", "xbox": "xbox", "win": "windows"}[k]
    qrow(k, lab, f"q.prod27.{pid}.{{}}", (f"prod27.{pid}.fy2025", f"prod27.{pid}.fy2026"), src=R8 + f"｜q.prod27.{pid}.*")
qcheck("chk_dc", "檢查：Devices and Consumer 產品線合計 − 分部", lambda c: f"SUM({c}{QR['srch']}:{c}{QR['win']})-{c}{QR['dc']}")
qrow("segsum", "分部合計（營收）", qform=lambda c: f"{c}{QR['ai']}+{c}{QR['dc']}", tag="Derived")
Q[f"L{QR['segsum']}"] = f"=L{QR['ai']}+L{QR['dc']}"; Q[f"M{QR['segsum']}"] = f"=M{QR['ai']}+M{QR['dc']}"
qrow("rev", "合併營收（公司損益表）", "q.rev.{}", ("is.rev.fy2025", "is.rev.fy2026"), src="財報新聞稿／10-K｜q.rev.*、is.rev.*")
qcheck("chk_rev", "檢查：分部合計 − 合併營收", lambda c: f"{c}{QR['segsum']}-{c}{QR['rev']}")
qsec("2. 分部營業利益（Microsoft 無未分攤項：分部合計＝合併營業利益）")
qrow("ai_oi", "Agents and Infra 營業利益", "q.seg27.ai.oi.{}", ("seg27.ai.oi.fy2025", "seg27.ai.oi.fy2026"), src=R8 + "｜q.seg27.ai.oi.*")
qrow("dc_oi", "Devices and Consumer 營業利益", "q.seg27.dc.oi.{}", ("seg27.dc.oi.fy2025", "seg27.dc.oi.fy2026"), src=R8 + "｜q.seg27.dc.oi.*")
qrow("oi", "合併營業利益（公司損益表）", "q.oi.{}", ("is.oi.fy2025", "is.oi.fy2026"), src="財報新聞稿／10-K｜q.oi.*、is.oi.*")
qcheck("chk_oi", "檢查：分部營業利益合計 − 合併營業利益（未分攤＝0）", lambda c: f"{c}{QR['ai_oi']}+{c}{QR['dc_oi']}-{c}{QR['oi']}")
qsec("3. 合併損益主要列（年度值；淨利另列各季）")
qrow("ni", "淨利（GAAP，含 OpenAI 等持股損益）", "q.ni.{}", ("is.ni.fy2025", "is.ni.fy2026"), src="財報新聞稿／10-K｜q.ni.*、is.ni.*")
for k, lab, a in (("oie", "業外收支（other income (expense), net）", "is.oie"), ("pbt", "稅前淨利", "is.pretax"), ("tax", "所得稅", "is.tax"), ("oai_g", "  其中：OpenAI 等持股損益（公司非 GAAP 調節）", "is.openaiGain")):
    qrow(k, lab, aids=(f"{a}.fy2025", f"{a}.fy2026"), sumcols=False, src=f"10-K｜{a}.*")
qrow("sh", "稀釋加權股數（十億股）", aids=("is.sharesDilW.fy2025", "is.sharesDilW.fy2026"), sumcols=False, src="10-K｜is.sharesDilW.*")
r_ = qrow("eps", "稀釋 EPS（US$）", sumcols=False, fmt="0.00", src="10-K｜is.epsDil.*")
Q[f"L{r_}"] = F["is.epsDil.fy2025"]["value"]; Q[f"M{r_}"] = F["is.epsDil.fy2026"]["value"]; Q[f"L{r_}"].font = BLUE; Q[f"M{r_}"].font = BLUE
qcheck("chk_pl", "檢查：營業利益＋業外 − 稅前；稅前 − 稅 − 淨利（年度，應為 0）", lambda c: f"({c}{QR['oi']}+{c}{QR['oie']}-{c}{QR['pbt']})+({c}{QR['pbt']}-{c}{QR['tax']}-{c}{QR['ni']})" if c in "LM" else None)
qsec("4. 現金流量（各季為新聞稿單季值；資本支出為現金 PP&E）")
qrow("ocf", "營業現金流", "q.cf.ocf.{}", ("cf.ocf.fy2025", "cf.ocf.fy2026"), src="新聞稿現金流量表｜q.cf.ocf.*")
qrow("capex", "資本支出（現金 PP&E，正值）", "q.cf.capex.{}", ("cf.capex.fy2025", "cf.capex.fy2026"), neg=True, src="新聞稿現金流量表｜q.cf.capex.*")
qrow("fcf", "自由現金流（營業現金流 − 現金資本支出）", qform=lambda c: f"{c}{QR['ocf']}-{c}{QR['capex']}", tag="Derived")
Q[f"L{QR['fcf']}"] = f"=L{QR['ocf']}-L{QR['capex']}"; Q[f"M{QR['fcf']}"] = f"=M{QR['ocf']}-M{QR['capex']}"
qrow("da", "折舊、攤銷及其他（現金流量表）", "q.cf.da.{}", ("cf.da.fy2025", "cf.da.fy2026"), src="新聞稿現金流量表｜q.cf.da.*（Q1 由年初至今推得）")
qrow("sbc", "股票薪酬", "q.cf.sbc.{}", ("cf.sbc.fy2025", "cf.sbc.fy2026"), src="新聞稿現金流量表｜q.cf.sbc.*")
qrow("bb", "回購（正值）", "q.cf.buyback.{}", ("cf.buyback.fy2025", "cf.buyback.fy2026"), neg=True, src="新聞稿現金流量表｜q.cf.buyback.*")
qrow("div", "股利（正值）", "q.cf.div.{}", ("cf.div.fy2025", "cf.div.fy2026"), neg=True, src="新聞稿現金流量表｜q.cf.div.*")
r_ = qrow("capex_co", "資本支出（公司口徑：含融資租賃，法說值）", sumcols=False, tag="Interested-party", src="法說｜capexFL.*（FY25Q4 起；公司口徑≠現金＋融資租賃）")
for c, i in zip("EFGHI", ("fy25q4", "fy26q1", "fy26q2", "fy26q3", "fy26q4")):
    Q[f"{c}{r_}"] = F[f"capexFL.{i}"]["value"] / 1000; Q[f"{c}{r_}"].font = BLUE; Q[f"{c}{r_}"].number_format = "#,##0.000"
Q[f"K{r_}"] = f"=SUM(F{r_}:I{r_})"; Q[f"M{r_}"] = fb("capexFL.fy2026"); Q[f"M{r_}"].font = BLUE
qrow("flp", "融資租賃本金（正值）", aids=("lease.finPrincipal.fy2025", "lease.finPrincipal.fy2026"), sumcols=False, src="10-K 附註 13｜lease.finPrincipal.*")
qcheck("chk_y", "檢查：四季合計 − 公司年度值（以上各列絕對值合計）", lambda c: (f"SUMPRODUCT(ABS(J{QR['ai']}:J{QR['oi']}-L{QR['ai']}:L{QR['oi']}))+ABS(J{QR['ni']}-L{QR['ni']})+SUMPRODUCT(ABS(J{QR['ocf']}:J{QR['div']}-L{QR['ocf']}:L{QR['div']}))" if c == "J" else
                                                                (f"SUMPRODUCT(ABS(K{QR['ai']}:K{QR['oi']}-M{QR['ai']}:M{QR['oi']}))+ABS(K{QR['ni']}-M{QR['ni']})+SUMPRODUCT(ABS(K{QR['ocf']}:K{QR['div']}-M{QR['ocf']}:M{QR['div']}))+ABS(K{QR['capex_co']}-M{QR['capex_co']})" if c == "K" else None)))
qsec("5. 對照：舊三分部（FY2026 以前口徑；10-K／各季新聞稿）與 Azure 年增")
for k, lab in (("pbp", "Productivity and Business Processes（舊）"), ("ic", "Intelligent Cloud（舊）"), ("mpc", "More Personal Computing（舊）")):
    qrow(f"old_{k}", lab + " 營收", f"q.seg.{k}.rev.{{}}", (f"seg.{k}.rev.fy2025", f"seg.{k}.rev.fy2026"), src=f"10-K／新聞稿｜q.seg.{k}.rev.*")
qcheck("chk_old", "檢查：舊三分部合計 − 合併營收", lambda c: f"{c}{QR['old_pbp']}+{c}{QR['old_ic']}+{c}{QR['old_mpc']}-{c}{QR['rev']}")
r_ = qrow("azg", "Azure 年增（重編口徑，公司揭露）", sumcols=False, fmt="0%", tag="Interested-party", src="8-K 2026-09-02 Investor Metrics｜q.azure27.growth.*")
for c, i in zip("FGHI", ("fy26q1", "fy26q2", "fy26q3", "fy26q4")):
    Q[f"{c}{r_}"] = F[f"q.azure27.growth.{i}"]["value"]; Q[f"{c}{r_}"].font = BLUE; Q[f"{c}{r_}"].number_format = "0%"
Q.column_dimensions["A"].width = 58
for c in "BCDEFGHIJKLM": Q.column_dimensions[c].width = 13
Q.column_dimensions["N"].width = 16; Q.column_dimensions["O"].width = 70
Q.freeze_panes = "B4"
# In 頁 sbc_r 公式回填 Q 列號
IN.cell(IR["sbc_r"], 5, f"=Q!M{QR['sbc']}/Q!M{QR['rev']}").number_format = "0.0%"


def qa(key, y):
    """Q 頁年度值（公司揭露）：FY2025＝L、FY2026＝M"""
    return f"Q!{'L' if y == 2025 else 'M'}{QR[key]}"


# =====================================================================================
# 引擎：一個情境的完整推導（Thesis／Seg／PL／CF／Val 各部分），主模型與 Sens 共用
# =====================================================================================
CUR = {"Thesis": 4, "Seg": 4, "PL": 12, "CF": 14, "Val": 70, "Sens": 70}


class Block:
    def __init__(self, s, title, route, P=None, V=None, base=None):
        self.s = s; self.title = title; self.route = route
        self.P = dict(pace="pace", cap=None, ls=f"ls_{s}", kmult=None, ext="ext_sh", life="life_it", g_azo="glt_azo", cop_add=f"cop_add_{s}",
                      cop_px="cop_px", g_srch="glt_srch", g_xbox="glt_xbox", dm="dm_seg", chip="chip_f", reading="internal_ai_reading")
        self.P.update(P or {})
        self.P = {k_: (v_.replace("{s}", s) if isinstance(v_, str) else v_) for k_, v_ in self.P.items()}
        self.V = dict(wacc="wacc", g="g_tv", ronic="ronic_prem", mm=None, m_sw="m_sw", m_cloud="m_cloud", m_ads="m_ads", m_game="m_game", m_pc="m_pc", m_ai="m_ai", oai="oai_v")
        self.V.update(V or {})
        self.R = dict(base.R) if base else {}
        self.specs = []

    def rng(self, key, y1, y2):
        """同一列的年度範圍：工作表名只寫一次（避免「CF!F30:CF!J30」被讀成 F 欄到 CF 欄）"""
        sh, r_, _ = self.R[key]
        return f"{sh}!{col(y1)}{r_}:{col(y2)}{r_}"

    def ref(self, key, y=None, absolute=False):
        sh, r_, scal = self.R[key]
        if scal:
            return f"{sh}!$C${r_}"
        c = col(y)
        return f"{sh}!${c}${r_}" if absolute else f"{sh}!{c}{r_}"

    def row(self, part, key, label, unit, f, fmt="#,##0.0", note="", years=YRS, scalar=False, check=False):
        self.specs.append(dict(part=part, key=key, label=label, unit=unit, f=f, fmt=fmt, note=note, years=years, scalar=scalar, check=check))

    def layout_and_write(self):
        started = set()
        pos = []
        for sp in self.specs:                              # 第一輪：配置列號
            sh = self.route[sp["part"]]
            if (sp["part"], sh) not in started:
                started.add((sp["part"], sh))
                pos.append(("title", sh, CUR[sh], sp["part"])); CUR[sh] += 1
                pos.append(("hdr", sh, CUR[sh], sp["part"])); CUR[sh] += 1
            self.R[sp["key"]] = (sh, CUR[sh], sp["scalar"])
            pos.append(("row", sh, CUR[sh], sp)); CUR[sh] += 1
        PARTN = {"thesis": "命題模組（Thesis）", "seg": "營收分部層（Seg）", "pl": "損益（PL）", "cf": "現金流與資金（CF）", "val": "評價（Val）"}
        for kind, sh, r_, x in pos:                         # 第二輪：寫入
            ws = WS[sh]
            if kind == "title":
                ws.cell(r_, 1, f"{self.title}｜{PARTN[x]}").font = BOLD
                continue
            if kind == "hdr":
                hdr(ws, r_, YHDR)
                continue
            sp = x
            ws.cell(r_, 1, sp["label"]); ws.cell(r_, 2, sp["unit"])
            if sp["scalar"]:
                v = sp["f"](None)
                c = ws.cell(r_, 3, ("=" + v) if isinstance(v, str) else v); c.number_format = sp["fmt"]
            else:
                for y in sp["years"]:
                    v = sp["f"](y)
                    if v is None:
                        continue
                    c = ws[f"{col(y)}{r_}"]
                    c.value = ("=" + v) if isinstance(v, str) else v
                    c.number_format = sp["fmt"]
            ws.cell(r_, 12, sp["note"])
            if sp["check"]:
                for cc in range(1, 12): ws.cell(r_, cc).fill = CHK


def gpath(b, key, glt):
    """近一年年增（FY2026 對 FY2025）線性收斂到長期值：FY(2026+i) 的年增＝g0 ＋（gLT − g0）× MIN(1, i ÷ n_conv)"""
    return lambda y: (f"{b.ref(key, 2026)}/{b.ref(key, 2025)}-1" if y == 2026 else
                      f"{b.ref('g_' + key, 2026, True)}+({glt}-{b.ref('g_' + key, 2026, True)})*MIN(1,({b.ref('yr', y)}-base_fy)/n_conv)" if y >= 2027 else None)


def dep_formula(b, capkey, y, life):
    """半年慣例直線折舊：第 0 年 0.5、第 1…L−1 年 1、第 L 年 0.5（L 可為變數）"""
    terms = []
    for v in range(2024, y + 1):
        d = y - v
        w = "0.5" if d == 0 else f"IF({d}<{life},1,IF({d}={life},0.5,0))"
        terms.append(f"{b.ref(capkey, v)}*{w}")
    return "(" + "+".join(terms) + f")/{life}"


def engine(b, val=True, ops=True):
    s = b.s; P = b.P; V = b.V; R = b.ref; RG = b.rng
    km = f"*{P['kmult']}" if P["kmult"] else ""
    sw = f"IF({P['reading']}=2,1,0)"
    FY26P = [y for y in YRS if y >= 2026]
    if ops:
        # ---------------- Thesis（命題模組）----------------
        t = "thesis"
        b.row(t, "yr", "會計年度（數值；公式以「年度 − base_fy」計年數）", "年", lambda y: y, fmt="0")
        pace_y = lambda y: (P["pace"] if y <= 2029 else
                            f"{P['pace']}+(g_tv*{R('ext', 2029)}-{P['pace']})*MIN(1,({R('yr', y)}-pace_y0)/(pace_y1-pace_y0))")
        b.row(t, "pace", "在役對外 AI MW 每年新增", "MW", lambda y: pace_y(y) if y >= 2027 else None, fmt="#,##0",
              note="FY2029 以前＝依資本支出指引的速度；之後逐年線性收斂，FY2032 只剩 g_tv × FY2029 在役 MW（修正一 F5：終值不落在建置高峰）")
        b.row(t, "ext", "在役對外 AI MW（年底，含租用）", "MW",
              lambda y: f"mw_ext_{y % 100}" if y <= 2026 else (f"MIN({P['cap']},{R('ext', y - 1)}+{R('pace', y)})" if P["cap"] else f"{R('ext', y - 1)}+{R('pace', y)}"),
              fmt="#,##0", note="三情境速度相同（變現驅動不同）")
        b.row(t, "rent", "  其中租用 neocloud 的對外 MW", "MW", lambda y: f"rent_ext_{y % 100}" if y <= 2026 else "rent_ext_f", fmt="#,##0", note="收入照算，不計資本支出與折舊（C20）")
        b.row(t, "own_ext", "  自有對外 MW", "MW", lambda y: f"{R('ext', y)}-{R('rent', y)}", fmt="#,##0")
        b.row(t, "own", "自有 AI MW（對外 ÷ 對外比例）", "MW", lambda y: f"{R('own_ext', y)}/{P['ext']}", fmt="#,##0", note="含自用（M365／GitHub Copilot、MAI 研發）")
        b.row(t, "inner", "  其中自用 AI MW", "MW", lambda y: f"{R('own', y)}-{R('own_ext', y)}", fmt="#,##0")
        b.row(t, "add_own", "自有 AI MW 新增（FY2024 欄含以前全部）", "MW", lambda y: R("own", y) if y == 2024 else f"{R('own', y)}-{R('own', y - 1)}", fmt="#,##0")
        b.row(t, "add_ext", "在役對外 MW 新增", "MW", lambda y: f"{R('ext', y)}-{R('ext', y - 1)}" if y >= 2025 else None, fmt="#,##0")
        ag3 = lambda y: f"a_g3_{min(y, 2029) % 100}"
        b.row(t, "rx", "  其中汰換為新世代的對外 MW（汰換 MW × 對外比例）", "MW", lambda y: f"{R('refresh', y)}*{P['ext']}" if y >= 2027 else None, fmt="#,##0",
              note="修正一：汰換時在役世代占比同步更新（舊世代依期初占比移出、新世代依當年新增占比移入）")
        b.row(t, "mw_h", "  在役對外 MW：Hopper", "MW", lambda y: f"mix_h*{R('ext', y)}" if y <= 2026 else f"{R('mw_h', y - 1)}-{R('rx', y)}*mix_h", fmt="#,##0")
        b.row(t, "mw_g2", "  在役對外 MW：GB200", "MW", lambda y: f"mix_g2*{R('ext', y)}" if y <= 2026 else f"{R('mw_g2', y - 1)}-{R('rx', y)}*mix_g2", fmt="#,##0")
        b.row(t, "mw_g3", "  在役對外 MW：GB300（含 Maia 200）", "MW",
              lambda y: f"mix_g3*{R('ext', y)}" if y <= 2026 else f"{R('mw_g3', y - 1)}+({R('add_ext', y)}+{R('rx', y)})*{ag3(y)}-{R('rx', y)}*mix_g3", fmt="#,##0")
        b.row(t, "mw_vr", "  在役對外 MW：VR200", "MW", lambda y: 0 if y <= 2026 else f"{R('mw_vr', y - 1)}+({R('add_ext', y)}+{R('rx', y)})*(1-{ag3(y)})", fmt="#,##0")
        b.row(t, "hold_w", "在役世代加權 IF_HoldEcon", "US$M/MW-年",
              lambda y: f"({R('mw_h', y)}*tk_hold_h100+{R('mw_g2', y)}*tk_hold_gb200+{R('mw_g3', y)}*tk_hold_gb300+{R('mw_vr', y)}*tk_hold_vr200)/{R('ext', y)}", fmt="0.000")
        b.row(t, "k", "k（長約占比 × k_長約＋其餘 × k_現貨）", "x",
              lambda y: "k_B" if y <= 2026 else f"(k_B+(({P['ls']}*k_long+(1-{P['ls']})*k_spot)-k_B)*MIN(1,({R('yr', y)}-base_fy)/k_ramp)){km}", fmt="0.000",
              note="實績期三情境同用基準 k（拆分是估計，不是情境）；FY2027 起 k_ramp 年線性移到情境 k")
        b.row(t, "rev_mw", "每 MW 年收入（對外）", "US$M/MW-年", lambda y: f"{R('hold_w', y)}*{R('k', y)}*{P['chip']}", fmt="0.000", note="按 MW-year 計價，不乘利用率（v0.2.2 口徑）")
        b.row(t, "azai", "Azure AI 營收（期間平均在役對外 MW × 每 MW 年收入）", "US$B", lambda y: f"({R('ext', y - 1)}+{R('ext', y)})/2*{R('rev_mw', y)}/1000" if y >= 2025 else None, note="FY2025、FY2026 為模型估計（用於拆出 Azure 其他）")
        b.row(t, "opex_w", "在役世代加權 IF_OpexGW", "US$M/MW-年",
              lambda y: f"({R('mw_h', y)}*tk_opex_h100+{R('mw_g2', y)}*tk_opex_gb200+{R('mw_g3', y)}*tk_opex_gb300+{R('mw_vr', y)}*tk_opex_vr200)/{R('ext', y)}", fmt="0.000")
        b.row(t, "c_own", "AI 營運成本：自有 AI MW（平均 × IF_OpexGW）", "US$B", lambda y: f"({R('own', y - 1)}+{R('own', y)})/2*{R('opex_w', y)}/1000" if y >= 2025 else None)
        b.row(t, "c_rent", "AI 營運成本：租用 neocloud 租金", "US$B", lambda y: f"({R('rent', y - 1)}+{R('rent', y)})/2*rent_mw/1000" if y >= 2025 else None)
        b.row(t, "c_fac", "AI 營運成本：租用機房租金（1 − 自建比例）", "US$B", lambda y: f"({R('own', y - 1)}+{R('own', y)})/2*(1-self_build)*fac_rent/1000" if y >= 2025 else None, note="C22")
        b.row(t, "ext_frac", "自有 AI 中對外的比重（期間平均）", "%", lambda y: f"({R('own_ext', y - 1)}+{R('own_ext', y)})/({R('own', y - 1)}+{R('own', y)})" if y >= 2025 else None, fmt="0.0%")
        b.row(t, "azai_ebitda", "Azure AI EBITDA（營收 − 對外營運成本 − 租金）", "US$B",
              lambda y: f"{R('azai', y)}-{R('ext_frac', y)}*({R('c_own', y)}+{R('c_fac', y)})-{R('c_rent', y)}" if y >= 2025 else None)
        b.row(t, "int_cost", "自用 AI 營運成本（服務 Copilot 等自有產品）", "US$B", lambda y: f"(1-{R('ext_frac', y)})*({R('c_own', y)}+{R('c_fac', y)})" if y >= 2025 else None)
        cx_add = lambda y: "cxit_open" if y <= 2026 else f"({ag3(y)}*tk_cxit_gb300+(1-{ag3(y)})*tk_cxit_vr200)"
        b.row(t, "cxit_add", "新增與汰換 MW 的 IT 資本支出（世代加權）", "US$M/MW", lambda y: cx_add(y), fmt="0.00", note="FY2026 以前用期初世代加權（估計）")
        b.row(t, "refresh", "GPU 汰換 MW（自有 AI MW 滿 IT 年限者）", "MW",
              lambda y: "+".join(f"{R('add_own', v)}*IF({R('yr', y)}-{R('yr', v)}={P['life']},1,0)" for v in range(2024, y)) if y >= 2027 else None, fmt="#,##0",
              note="只換 IT、機房不換；FY2024 欄含 FY2024 以前全部（視同 FY2024 投入）；舊機隊在年限到期後停止折舊，汰換後的新 IT 另計折舊（不重複）")
        b.row(t, "cap_it0", "AI IT 資本支出（公式）", "US$B", lambda y: f"({R('add_own', y)}+{R('refresh', y)})*{R('cxit_add', y)}/1000" if y >= 2027 else f"{R('add_own', y)}*{R('cxit_add', y)}/1000")
        b.row(t, "cap_fac", "AI 機房資本支出（自建比例 × IF_CapexFacility）", "US$B", lambda y: f"{R('add_own', y)}*self_build*tk_cxfac/1000")
        b.row(t, "floor_x", "FY2027 補足到公司資本支出下限的差額（計入 IT）", "US$B",
              lambda y: f"MAX(0,capex_floor27-{R('cap_it0', y)}-{R('cap_fac', y)}-{R('nacx', y)})" if y == 2027 else None,
              note="近期以公司指引為下限（C10 方法沿用）；IT 占約三分之二（法說短期資產）")
        b.row(t, "cap_it", "AI IT 資本支出（含補足差額）", "US$B", lambda y: f"{R('cap_it0', y)}+{R('floor_x', y)}" if y == 2027 else R("cap_it0", y))
        b.row(t, "ai_capex", "AI 資本支出合計", "US$B", lambda y: f"{R('cap_it', y)}+{R('cap_fac', y)}", note="FY2026 以前為依 MW 估計的值（非公司揭露）")
        b.row(t, "dep_it", "AI IT 折舊（半年慣例、年限 life_it）", "US$B", lambda y: dep_formula(b, "cap_it", y, P["life"]) if y >= 2026 else None)
        b.row(t, "dep_fac", "AI 機房折舊（年限 life_fac）", "US$B", lambda y: dep_formula(b, "cap_fac", y, "life_fac") if y >= 2026 else None)
        b.row(t, "ai_da", "AI 折舊合計（進 Agents and Infra 營業利益）", "US$B", lambda y: f"{R('dep_it', y)}+{R('dep_fac', y)}" if y >= 2026 else None)
        b.row(t, "azai_ebit", "Azure AI 營業利益（EBITDA − 對外分攤折舊）", "US$B", lambda y: f"{R('azai_ebitda', y)}-{R('ext_frac', y)}*{R('ai_da', y)}" if y >= 2026 else None)
        # ---------------- Seg（營收分部層）----------------
        g = "seg"
        SY = [y for y in YRS if y >= 2025]
        b.row(g, "s_azai", "  Azure AI（命題子列）", "US$B", lambda y: R("azai", y), years=SY)
        b.row(g, "azo", "  Azure 其他（非 AI 雲端；FY2025–26＝Azure 實績 − Azure AI 估計）", "US$B",
              lambda y: f"{qa('azure', y)}-{R('s_azai', y)}" if y <= 2026 else f"{R('azo', y - 1)}*(1+{R('g_azo', y)})", years=SY)
        b.row(g, "g_azo", "    年增（近一年 → 長期值）", "%", gpath(b, "azo", P["g_azo"]), fmt="0.0%", years=FY26P)
        b.row(g, "azure", "Azure", "US$B", lambda y: f"{R('s_azai', y)}+{R('azo', y)}", years=SY)
        b.row(g, "seats", "  Copilot 付費座席（年底）", "百萬席",
              lambda y: f"cop_s{y % 100}" if y <= 2026 else f"{R('seats', y - 1)}+{P['cop_add']}" + ("" if y <= 2029 else "*cop_tail"), fmt="#,##0",
              note="FY2027–FY2029 每年淨增 cop_add；FY2030 起 × cop_tail")
        b.row(g, "cop", "  Copilot（命題子列：平均座席 × 月費 × 12）", "US$B", lambda y: f"({R('seats', y - 1)}+{R('seats', y)})/2*{'cop_px' if y <= 2026 else P['cop_px']}*12/1000", years=SY,
              note="實績期用基準月費（拆分是估計）；敏感度只改預測期")
        b.row(g, "m365b", "  M365 基礎（不含 Copilot；FY2025–26＝實績 − Copilot 估計）", "US$B",
              lambda y: f"{qa('m365', y)}-{R('cop', y)}" if y <= 2026 else f"{R('m365b', y - 1)}*(1+{R('g_m365b', y)})", years=SY)
        b.row(g, "g_m365b", "    年增（近一年 → 長期值）", "%", gpath(b, "m365b", "glt_m365b"), fmt="0.0%", years=FY26P)
        b.row(g, "m365", "Microsoft 365 cloud", "US$B", lambda y: f"{R('m365b', y)}+{R('cop', y)}", years=SY)
        for k, lab, glt in (("lic", "Productivity and server licensing", "glt_lic"), ("ind", "Industry solutions", "glt_ind"), ("fro", "Frontier and support services", "glt_fro")):
            b.row(g, k, lab, "US$B", (lambda k_: lambda y: qa(k_, y) if y <= 2026 else f"{R(k_, y - 1)}*(1+{R('g_' + k_, y)})")(k), years=SY)
            b.row(g, "g_" + k, "    年增（近一年 → 長期值）", "%", gpath(b, k, glt), fmt="0.0%", years=FY26P)
        b.row(g, "ai_seg", "Agents and Infra", "US$B", lambda y: f"{R('azure', y)}+{R('m365', y)}+{R('lic', y)}+{R('ind', y)}+{R('fro', y)}", years=SY)
        for k, lab, glt in (("srch", "Search and advertising", P["g_srch"]), ("xbox", "XBOX", P["g_xbox"]), ("win", "Windows OEM and devices", "glt_win")):
            b.row(g, k, lab, "US$B", (lambda k_: lambda y: qa(k_, y) if y <= 2026 else f"{R(k_, y - 1)}*(1+{R('g_' + k_, y)})")(k), years=SY)
            b.row(g, "g_" + k, "    年增（近一年 → 長期值）", "%", gpath(b, k, glt), fmt="0.0%", years=FY26P)
        b.row(g, "dc", "Devices and Consumer", "US$B", lambda y: f"{R('srch', y)}+{R('xbox', y)}+{R('win', y)}", years=SY)
        b.row(g, "tot", "合併營收（分部合計）", "US$B", lambda y: f"{R('ai_seg', y)}+{R('dc', y)}", years=SY)
        b.row(g, "c_seg", "檢查：產品線合計 − 分部；實績期另比對公司分部值", "US$B",
              lambda y: (f"({R('azure', y)}+{R('m365', y)}+{R('lic', y)}+{R('ind', y)}+{R('fro', y)}-{R('ai_seg', y)})+({R('srch', y)}+{R('xbox', y)}+{R('win', y)}-{R('dc', y)})"
                         + (f"+({R('ai_seg', y)}-{qa('ai', y)})+({R('dc', y)}-{qa('dc', y)})+({R('azure', y)}-{qa('azure', y)})+({R('m365', y)}-{qa('m365', y)})" if y <= 2026 else "")),
              fmt="0.000", years=SY, check=True)
        # ---------------- PL ----------------
        p = "pl"
        PY = SY
        for k, lab in (("azure", "Azure"), ("s_azai", "  Azure AI（命題子列）"), ("azo", "  Azure 其他"), ("m365", "Microsoft 365 cloud"), ("m365b", "  M365 基礎（不含 Copilot）"),
                       ("cop", "  Copilot（命題子列）"), ("lic", "Productivity and server licensing"), ("ind", "Industry solutions"), ("fro", "Frontier and support services"),
                       ("ai_seg", "Agents and Infra"), ("srch", "Search and advertising"), ("xbox", "XBOX"), ("win", "Windows OEM and devices"), ("dc", "Devices and Consumer"), ("tot", "合併營收")):
            b.row(p, "p_" + k, lab, "US$B", (lambda k_: lambda y: R(k_, y))(k), years=PY)
        b.row(p, "p_g", "營收年增", "%", lambda y: f"{R('p_tot', y)}/{R('p_tot', y - 1)}-1", fmt="0.0%", years=FY26P)
        b.row(p, "p_chk", "檢查：分部合計 − 合併營收；子列合計 − 母列（Azure、M365 cloud）", "US$B",
              lambda y: f"({R('p_ai_seg', y)}+{R('p_dc', y)}-{R('p_tot', y)})+({R('p_s_azai', y)}+{R('p_azo', y)}-{R('p_azure', y)})+({R('p_m365b', y)}+{R('p_cop', y)}-{R('p_m365', y)})"
                        + (f"+({R('p_tot', y)}-{qa('rev', y)})" if y <= 2026 else ""), fmt="0.000", years=PY, check=True)
        b.row(p, "m0", "Agents and Infra 核心利益率（扣 AI 營運成本與 AI 折舊前；FY2026 校準）", "%",
              lambda y: (f"({qa('ai_oi', 2026)}+{R('c_own', 2026)}+{R('c_rent', 2026)}+{R('c_fac', 2026)}+{R('ai_da', 2026)}-{R('azai', 2026)})/({qa('ai', 2026)}-{R('azai', 2026)})" if y == 2026
                         else f"{R('m0', 2026, True)}+{P['dm']}"), fmt="0.0%", years=FY26P,
              note="FY2026：(分部營業利益＋AI 營運成本＋AI 折舊 − Azure AI 營收) ÷ (分部營收 − Azure AI 營收)；之後固定（＋dm_seg）")
        b.row(p, "ai_oi", "Agents and Infra 營業利益", "US$B",
              lambda y: qa("ai_oi", y) if y <= 2026 else f"{R('m0', y)}*({R('ai_seg', y)}-{R('azai', y)})+{R('azai', y)}-({R('c_own', y)}+{R('c_rent', y)}+{R('c_fac', y)})-{R('ai_da', y)}",
              years=PY, note="＝核心利益率 ×（分部營收 − Azure AI）＋ Azure AI 營收 − AI 營運成本 − AI 折舊：AI 資本支出經折舊進分部營業利益")
        b.row(p, "m_dc", "Devices and Consumer 營業利益率", "%", lambda y: f"{qa('dc_oi', 2026)}/{qa('dc', 2026)}" if y == 2026 else f"{R('m_dc', 2026, True)}+{P['dm']}",
              fmt="0.0%", years=FY26P, note="FY2026 實績後固定（＋dm_seg）")
        b.row(p, "dc_oi", "Devices and Consumer 營業利益", "US$B", lambda y: qa("dc_oi", y) if y <= 2026 else f"{R('dc', y)}*{R('m_dc', y)}", years=PY)
        b.row(p, "oi", "合併營業利益（財報口徑；分部合計，無未分攤項）", "US$B", lambda y: f"{R('ai_oi', y)}+{R('dc_oi', y)}", years=PY)
        b.row(p, "oi_chk", "檢查：分部營業利益合計 − 公司合併營業利益（實績）", "US$B", lambda y: f"{R('oi', y)}-{qa('oi', y)}" if y <= 2026 else None, fmt="0.000", years=PY, check=True)
        b.row(p, "oim", "營業利益率（財報口徑）", "%", lambda y: f"{R('oi', y)}/{R('p_tot', y)}", fmt="0.0%", years=PY)
        b.row(p, "p_shadow", "自用 AI 影子效益（讀法 2；非財報營收；分攤到 Agents and Infra）", "US$B", lambda y: R("shadow", y), years=FCY,
              note="修正一 F1：營收與分部對帳不動；只進評價口徑的營業利益、淨利、EPS 與 DCF")
        b.row(p, "oi_val", "營業利益（評價口徑＝財報＋自用 AI 影子效益）", "US$B", lambda y: R("oi", y) if y <= 2026 else f"{R('oi', y)}+{R('p_shadow', y)}", years=PY)
        b.row(p, "p_aida", "  備忘：AI 折舊（Thesis）", "US$B", lambda y: R("ai_da", y), years=FY26P)
        b.row(p, "p_nada", "  備忘：非 AI 折舊攤銷（CF）", "US$B", lambda y: R("nada", y), years=FY26P)
        b.row(p, "oie", "業外收支（預測期不含持股評價損益＝0）", "US$B", lambda y: qa("oie", y) if y <= 2026 else "oie_f", years=PY, note="FY2026 含 OpenAI 重組利得等 +6.5")
        b.row(p, "pbt", "稅前淨利（財報口徑）", "US$B", lambda y: f"{R('oi', y)}+{R('oie', y)}", years=PY)
        b.row(p, "tax", "所得稅（財報口徑）", "US$B", lambda y: qa("tax", y) if y <= 2026 else f"{R('pbt', y)}*tax_f", years=PY)
        b.row(p, "ni", "淨利（財報口徑）", "US$B", lambda y: f"{R('pbt', y)}-{R('tax', y)}", years=PY)
        b.row(p, "sh", "稀釋股數（十億股；回購依 CF 頁）", "十億股", lambda y: qa("sh", y) if y <= 2026 else R("dil", y), fmt="0.000", years=PY)
        b.row(p, "eps_r", "EPS（財報口徑，US$）", "US$", lambda y: f"{R('ni', y)}/{R('sh', y)}", fmt="0.00", years=PY, note="實績期與公司 EPS 差 < 0.01（四捨五入）")
        b.row(p, "ni_val", "淨利（評價口徑＝財報＋影子效益 ×（1 − 稅率））", "US$B", lambda y: R("ni", y) if y <= 2026 else f"{R('ni', y)}+{R('p_shadow', y)}*(1-tax_f)", years=PY)
        b.row(p, "eps", "EPS（評價口徑，US$）", "US$", lambda y: f"{R('ni_val', y)}/{R('sh', y)}", fmt="0.00", years=PY, note="讀法 2 主值；讀法 1 時＝財報口徑")
        b.row(p, "eps_g", "EPS 年增（評價口徑）", "%", lambda y: f"{R('eps', y)}/{R('eps', y - 1)}-1", fmt="0.0%", years=FY26P)
        # ---------------- Thesis 第二段：AI 增量報酬（需 Copilot 與核心利益率）----------------
        b.row(t, "t_cop", "Copilot 營收（Seg）", "US$B", lambda y: R("cop", y), years=FY26P)
        b.row(t, "shadow_full", "自用 AI 影子效益（讀法 2 金額）＝MAX(0, 平均自用 MW × 持有成本（k＝1）− Copilot × 核心利益率)", "US$B",
              lambda y: f"MAX(0,({R('inner', y - 1)}+{R('inner', y)})/2*{R('hold_w', y)}/1000-{R('t_cop', y)}*{R('m0', y)})", years=FCY,
              note="自用 AI 價值中性：自用算力至少回收 Tokenomics 持有成本；Copilot 已在 M365 營收內計入的部分扣除，避免重複")
        b.row(t, "shadow", "自用 AI 影子效益（依讀法開關 internal_ai_reading）", "US$B", lambda y: f"{sw}*{R('shadow_full', y)}", years=FCY)
        b.row(t, "ai_ebit", "AI 增量營業利益（評價口徑）＝Azure AI EBITDA＋Copilot × 核心利益率＋影子效益 − 自用 AI 營運成本 − AI 折舊", "US$B",
              lambda y: f"{R('azai_ebitda', y)}+{R('t_cop', y)}*{R('m0', y)}+{R('shadow', y)}-{R('int_cost', y)}-{R('ai_da', y)}", years=FCY)
        b.row(t, "nppe", "AI PP&E 淨額（年底）", "US$B",
              lambda y: f"{R('own', 2026)}*(cxit_open+self_build*tk_cxfac)/1000*ai_net_open" if y == 2026 else f"{R('nppe', y - 1)}+{R('ai_capex', y)}-{R('ai_da', y)}",
              years=FY26P, note="期初＝自有 AI MW × 每 MW 全成本 × 淨額比（v0.2.2）")
        b.row(t, "roic", "AI 增量 ROIC（評價口徑；稅後營業利益 ÷ 平均 AI PP&E 淨額）", "%", lambda y: f"{R('ai_ebit', y)}*(1-tax_f)/(({R('nppe', y - 1)}+{R('nppe', y)})/2)", fmt="0.0%", years=FCY)
        b.row(t, "roic_r1", "  讀法 1（自用 AI 只計成本）的 AI 增量 ROIC", "%",
              lambda y: f"({R('ai_ebit', y)}-{R('shadow', y)})*(1-tax_f)/(({R('nppe', y - 1)}+{R('nppe', y)})/2)", fmt="0.0%", years=FCY)
        b.row(t, "roic_r2", "  讀法 2（自用 AI 價值中性）的 AI 增量 ROIC", "%",
              lambda y: f"({R('ai_ebit', y)}-{R('shadow', y)}+{R('shadow_full', y)})*(1-tax_f)/(({R('nppe', y - 1)}+{R('nppe', y)})/2)", fmt="0.0%", years=FCY)
        b.row(t, "roic_x", "  對外 AI ROIC（Azure AI 營業利益 ÷ 對外分攤的 PP&E；C11 口徑）", "%",
              lambda y: f"{R('azai_ebit', y)}*(1-tax_f)/({R('ext_frac', y)}*({R('nppe', y - 1)}+{R('nppe', y)})/2)", fmt="0.0%", years=FCY)
        b.row(t, "spread", "AI 增量 ROIC − WACC", "pt", lambda y: f"{R('roic', y)}-{V['wacc']}", fmt="0.0%", years=FCY)
        b.row(t, "ep", "AI 經濟利潤＝稅後 AI 營業利益 − WACC × 平均 AI PP&E 淨額", "US$B",
              lambda y: f"{R('ai_ebit', y)}*(1-tax_f)-{V['wacc']}*({R('nppe', y - 1)}+{R('nppe', y)})/2", years=FCY)
        # ---------------- CF（財報現金；收支關卡）----------------
        c = "cf"
        b.row(c, "c_ni", "淨利（財報口徑）", "US$B", lambda y: R("ni", y), years=FCY)
        b.row(c, "nada", "非 AI 折舊攤銷", "US$B", lambda y: f"{qa('da', 2026)}-{R('ai_da', 2026)}" if y == 2026 else f"{R('nada', y - 1)}*(1-1/life_nonai)+{R('nacx', y)}/life_nonai",
              years=FY26P, note="FY2026＝公司 D&A − AI 折舊估計（C17）")
        b.row(c, "da", "折舊攤銷合計（AI＋非 AI）", "US$B", lambda y: f"{R('ai_da', y)}+{R('nada', y)}", years=FY26P)
        b.row(c, "sbc", "股票薪酬（現金流量表加回；DCF 視為現金成本）", "US$B", lambda y: f"{R('tot', y)}*sbc_r", years=FCY)
        b.row(c, "wc", "營運資金變動", "US$B", lambda y: "wc27" if y == 2027 else "wc_f", years=FCY, note="C14：首期＝上一年度實績；之後 wc_f")
        b.row(c, "ocf", "營業現金流", "US$B", lambda y: qa("ocf", 2026) if y == 2026 else f"{R('c_ni', y)}+{R('da', y)}+{R('sbc', y)}+{R('wc', y)}", years=FY26P)
        b.row(c, "aicx", "AI 資本支出（Thesis）", "US$B", lambda y: R("ai_capex", y), years=FCY)
        b.row(c, "nacx", "非 AI 資本支出（各線營收 × 比例）", "US$B",
              lambda y: f"{R('azo', y)}*cx_azo+{R('m365', y)}*cx_m365+{R('lic', y)}*cx_lic+({R('ind', y)}+{R('fro', y)})*cx_ind+{R('srch', y)}*cx_srch+({R('xbox', y)}+{R('win', y)})*cx_dev", years=FY26P)
        b.row(c, "capex", "資本支出合計（公司口徑，含融資租賃）", "US$B", lambda y: f"Q!M{QR['capex_co']}" if y == 2026 else f"{R('aicx', y)}+{R('nacx', y)}", years=FY26P,
              note="新的融資租賃在取得時視同現金（偏保守，v0.2.2 口徑）；FY2026 為公司法說值")
        b.row(c, "fl_bal", "在帳融資租賃負債（年底）", "US$B", lambda y: "fl_liab" if y == 2026 else f"{R('fl_bal', y - 1)}-{R('fl_prin', y)}", years=FY26P)
        b.row(c, "fl_prin", "在帳融資租賃本金（付款 − 利息）", "US$B", lambda y: f"fl_pay_{y % 100}-{R('fl_bal', y - 1)}*fl_rate", years=FCY)
        b.row(c, "fcf", "自由現金流（財報現金：營業現金流 − 資本支出 − 融資租賃本金）", "US$B",
              lambda y: f"{qa('ocf', 2026)}-{qa('capex', 2026)}-{qa('flp', 2026)}" if y == 2026 else f"{R('ocf', y)}-{R('capex', y)}-{R('fl_prin', y)}", years=FY26P,
              note="FY2026 為實績（現金資本支出口徑）；收支關卡用此列（不含影子效益）")
        b.row(c, "dps", "每股股利", "US$", lambda y: "dps_27" if y == 2027 else f"{R('dps', y - 1)}*(1+dps_g)", fmt="0.00", years=FCY)
        b.row(c, "div", "股利", "US$B", lambda y: qa("div", 2026) if y == 2026 else f"{R('dps', y)}*{R('sh_end', y - 1)}", years=FY26P, note="每股股利 × 期初流通股數（避免與回購循環）")
        b.row(c, "fcfd", "自由現金流扣股利", "US$B", lambda y: f"{R('fcf', y)}-{R('div', y)}", years=FY26P)
        b.row(c, "repay", "債務到期償還", "US$B", lambda y: f"debt_mat_{y % 100}", years=FCY)
        b.row(c, "bb", "回購（現金不足時先減）", "US$B", lambda y: qa("bb", 2026) if y == 2026 else
              f"MAX(0,MIN(bb_target,{R('cash', y - 1)}+{R('fcfd', y)}-{R('repay', y)}-min_cash))", years=FY26P)
        b.row(c, "nd", "新增舉債（維持最低現金）", "US$B", lambda y: f"MAX(0,min_cash-({R('cash', y - 1)}+{R('fcfd', y)}-{R('repay', y)}-{R('bb', y)}))", years=FCY)
        b.row(c, "cash", "現金與短期投資（年底）", "US$B", lambda y: "cash_open" if y == 2026 else f"{R('cash', y - 1)}+{R('fcfd', y)}-{R('repay', y)}-{R('bb', y)}+{R('nd', y)}", years=FY26P)
        b.row(c, "debt", "債務（年底，面額）", "US$B", lambda y: "debt_face" if y == 2026 else f"{R('debt', y - 1)}-{R('repay', y)}+{R('nd', y)}", years=FY26P)
        b.row(c, "netcash", "淨現金（不含租賃）", "US$B", lambda y: f"{R('cash', y)}-{R('debt', y)}", years=FY26P)
        b.row(c, "bbpx", "回購均價", "US$", lambda y: f"px*(1+bb_g)^({R('yr', y)}-base_fy-mid_yr)", fmt="#,##0", years=FCY)
        b.row(c, "sh_end", "流通股數（年底）", "十億股", lambda y: "sh_out26" if y == 2026 else f"{R('sh_end', y - 1)}+sh_iss-{R('bb', y)}/{R('bbpx', y)}", fmt="0.000", years=FY26P)
        b.row(c, "dil", "稀釋加權股數", "十億股", lambda y: "sh_dil26" if y == 2026 else f"({R('sh_end', y - 1)}+{R('sh_end', y)})/2+dil_gap", fmt="0.000", years=FY26P)
        b.row(c, "gate", "自由現金流扣股利（FY2027–FY2031 合計，本情境）", "US$B", lambda y: f"SUM({RG('fcfd', 2027, 2031)})", scalar=True)
        b.row(c, "gate_chk", "檢查：收支關卡＝首年起五年自由現金流扣股利（逐欄加總）", "US$B",
              lambda y: f"{R('gate')}-(" + "+".join(R("fcfd", yy) for yy in range(2027, 2032)) + ")", fmt="0.000", scalar=True, check=True)
        def need_f(y):
            nd5, bb5 = RG("nd", 2027, 2031), RG("bb", 2027, 2031)
            return (f'IF(SUM({nd5})>0.05,"需要舉債 "&TEXT(SUM({nd5}),"0.0")&" 十億美元","不需舉債")&"；"&'
                         f'IF(SUM({bb5})<COUNT({bb5})*bb_target-0.05,"回購較每年 "&TEXT(bb_target,"0.0")&" 的目標少 "&TEXT(COUNT({bb5})*bb_target-SUM({bb5}),"0.0")&" 十億美元（"&TEXT(COUNTIF({bb5},"<"&(bb_target-0.05)),"0")&" 年減少）","回購不減")')
        b.row(c, "need", "是否需要舉債或停止回購（FY2027–FY2031）", "文字", need_f, fmt="@", scalar=True)
    if val:
        v = "val"
        W = V["wacc"]
        mm = f"*{V['mm']}" if V["mm"] else ""
        b.row(v, "tm", "折現年數（期中，自評價日）", "年", lambda y: f"{R('yr', y)}-base_fy-mid_yr", fmt="0.0", years=FCY)
        b.row(v, "df", "折現因子", "x", lambda y: f"1/(1+{W})^{R('tm', y)}", fmt="0.000", years=FCY)
        b.row(v, "fcf_val", "評價用自由現金流＝營業現金流 − 股票薪酬 − 資本支出＋影子效益 ×（1 − 稅率）", "US$B",
              lambda y: f"{R('ocf', y)}-{R('sbc', y)}-{R('capex', y)}+{R('shadow', y)}*(1-tax_f)", years=FCY,
              note="修正一：F2 股票薪酬視為現金成本；F1 讀法 2 影子效益；在帳融資租賃本金不扣，改在權益橋減負債（F6）")
        b.row(v, "pv", "評價用自由現金流現值", "US$B", lambda y: f"{R('fcf_val', y)}*{R('df', y)}", years=FCY)
        b.row(v, "sum_pv", "明確期現值合計（FY2027–FY2032）", "US$B", lambda y: f"SUM({RG('pv', 2027, 2032)})", scalar=True)
        b.row(v, "nopat", "永續起點 NOPAT（FY2032 評價口徑營業利益 ×（1＋g）×（1 − 稅率））", "US$B", lambda y: f"{R('oi_val', 2032)}*(1+{V['g']})*(1-tax_f)", scalar=True)
        b.row(v, "tvfcf", "永續期自由現金流（NOPAT ×（1 − g ÷ 新投資報酬））", "US$B", lambda y: f"{R('nopat')}*(1-{V['g']}/({W}+{V['ronic']}))", scalar=True)
        b.row(v, "tv", "永續價值（FY2032 年底）", "US$B", lambda y: f"{R('tvfcf')}/({W}-{V['g']})", scalar=True)
        b.row(v, "pvtv", "永續價值現值", "US$B", lambda y: f"{R('tv')}/(1+{W})^({R('tm', 2032)}+mid_yr)", scalar=True)
        b.row(v, "ev", "DCF 企業價值（2026-06-30）", "US$B", lambda y: f"{R('sum_pv')}+{R('pvtv')}", scalar=True)
        b.row(v, "tvs", "永續價值占比", "%", lambda y: f"{R('pvtv')}/{R('ev')}", fmt="0.0%", scalar=True)
        b.row(v, "eq", "股東權益價值＝EV＋淨現金 − 在帳融資租賃負債（FY2026 年底）＋OpenAI 持股", "US$B",
              lambda y: f"{R('ev')}+{R('netcash', 2026)}-{R('fl_bal', 2026)}+{V['oai']}", scalar=True)
        b.row(v, "ps0", "DCF 每股價值（2026-06-30）", "US$", lambda y: f"{R('eq')}/sh_dil26", fmt="#,##0.00", scalar=True)
        b.row(v, "divroll", "評價日到目標價時點已付的每股股利", "US$", lambda y: f"dps_27+{R('dps', 2028)}*roll_div_q/q_per_yr", fmt="0.00", scalar=True)
        b.row(v, "dcf12", "DCF 每股價值推到 12 個月後（×（1＋WACC）^t_roll − 期間股利）", "US$", lambda y: f"{R('ps0')}*(1+{W})^t_roll-{R('divroll')}", fmt="#,##0.00", scalar=True)
        b.row(v, "dar", "非 AI 折舊 ÷（營收 − Azure AI）", "%", lambda y: f"{R('nada', y)}/({R('tot', y)}-{R('azai', y)})", fmt="0.0%", years=FCY)
        LINES = (("e_m365", "Microsoft 365 cloud", "m365", "al_m365"), ("e_lic", "Productivity and server licensing", "lic", "al_lic"), ("e_ind", "Industry solutions", "ind", "al_ind"),
                 ("e_fro", "Frontier and support services", "fro", "al_fro"), ("e_srch", "Search and advertising", "srch", "al_srch"))
        for ek, lab, rk, al in LINES:
            b.row(v, ek, f"EBITDA：{lab}（營收 ×（分配利益率＋折舊率＋股票薪酬率））", "US$B", (lambda rk_, al_: lambda y: f"{R(rk_, y)}*({al_}+{R('dar', y)}+sbc_r)")(rk, al), years=FCY,
                  note="F2：同業 EV/EBITDA 多為股票薪酬加回口徑，分部 EBITDA 同口徑" if ek == "e_m365" else "")
        b.row(v, "e_aimod", "EBITDA：AI 基礎設施模組（Azure AI EBITDA − 自用 AI 營運成本＋影子效益）", "US$B",
              lambda y: f"{R('azai_ebitda', y)}-{R('int_cost', y)}+{R('shadow', y)}", years=FCY)
        b.row(v, "e_aiseg", "EBITDA：Agents and Infra 合計（評價口徑營業利益＋AI 折舊＋非 AI 折舊與股票薪酬分攤）", "US$B",
              lambda y: f"{R('ai_oi', y)}+{R('shadow', y)}+{R('ai_da', y)}+{R('dar', y)}*({R('ai_seg', y)}-{R('azai', y)})+sbc_r*{R('ai_seg', y)}", years=FCY)
        b.row(v, "e_azo", "EBITDA：Azure 其他（分部殘差；不含 AI 模組）", "US$B",
              lambda y: f"{R('e_aiseg', y)}-{R('e_m365', y)}-{R('e_lic', y)}-{R('e_ind', y)}-{R('e_fro', y)}-{R('e_aimod', y)}", years=FCY)
        b.row(v, "e_dc", "EBITDA：Devices and Consumer 合計", "US$B", lambda y: f"{R('dc_oi', y)}+{R('dar', y)}*{R('dc', y)}+sbc_r*{R('dc', y)}", years=FCY)
        b.row(v, "e_xbox", "EBITDA：XBOX（分部殘差按營收分）", "US$B", lambda y: f"({R('e_dc', y)}-{R('e_srch', y)})*{R('xbox', y)}/({R('xbox', y)}+{R('win', y)})", years=FCY)
        b.row(v, "e_win", "EBITDA：Windows OEM and devices", "US$B", lambda y: f"{R('e_dc', y)}-{R('e_srch', y)}-{R('e_xbox', y)}", years=FCY)
        # F3：AI 模組自身 DCF（同一 WACC；FCF＝AI EBITDA − 稅 − AI 資本支出〔含汰換〕）
        b.row(v, "ai_fcf", "AI 模組自由現金流＝EBITDA − 稅 ×（EBITDA − AI 折舊）− AI 資本支出（含汰換）", "US$B",
              lambda y: f"{R('e_aimod', y)}-tax_f*({R('e_aimod', y)}-{R('ai_da', y)})-{R('ai_capex', y)}", years=FCY)
        b.row(v, "ai_pv", "AI 模組自由現金流現值", "US$B", lambda y: f"{R('ai_fcf', y)}*{R('df', y)}", years=FCY)
        b.row(v, "ai_tv", "AI 模組永續價值現值（NOPAT ×（1 − g ÷ 新投資報酬）÷（WACC − g））", "US$B",
              lambda y: f"({R('e_aimod', 2032)}-{R('ai_da', 2032)})*(1-tax_f)*(1+{V['g']})*(1-{V['g']}/({W}+{V['ronic']}))/({W}-{V['g']})/(1+{W})^({R('tm', 2032)}+mid_yr)", scalar=True)
        b.row(v, "ne_aimod", "分部加總：AI 基礎設施模組（模組 DCF，推到目標價時點）", "US$B",
              lambda y: f"(SUM({RG('ai_pv', 2027, 2032)})+{R('ai_tv')})*(1+{W})^t_roll", scalar=True, note="F3：AI 腿以自身 DCF 計入資本密集度")
        b.row(v, "ne_azai_cmp", "對照：Azure AI NTM EBITDA × 同業倍數（v0.3 初版口徑）", "US$B",
              lambda y: f"(w_ntm28*{R('azai_ebitda', 2028)}+(1-w_ntm28)*{R('azai_ebitda', 2029)})*{V['m_ai']}{mm}", scalar=True)
        SOTP = (("e_azo", "Azure 其他", V["m_cloud"]), ("e_m365", "Microsoft 365 cloud", V["m_sw"]), ("e_lic", "Productivity and server licensing", V["m_sw"]),
                ("e_ind", "Industry solutions", V["m_sw"]), ("e_fro", "Frontier and support services", V["m_sw"]), ("e_srch", "Search and advertising", V["m_ads"]),
                ("e_xbox", "XBOX", V["m_game"]), ("e_win", "Windows OEM and devices", V["m_pc"]))
        for ek, lab, m in SOTP:
            b.row(v, "n" + ek, f"分部加總：{lab}（NTM EBITDA × 倍數）", "US$B",
                  (lambda ek_, m_: lambda y: f"(w_ntm28*{R(ek_, 2028)}+(1-w_ntm28)*{R(ek_, 2029)})*{m_}{mm}")(ek, m), scalar=True)
        b.row(v, "sev", "分部加總 EV（目標價時點）", "US$B", lambda y: R("ne_aimod") + "+" + "+".join(R("n" + ek) for ek, _, _ in SOTP), scalar=True)
        b.row(v, "sbr", "加：淨現金（FY2027 年底）＋OpenAI 持股 − 在帳融資租賃負債", "US$B", lambda y: f"{R('netcash', 2027)}+{V['oai']}-{R('fl_bal', 2027)}", scalar=True,
              note="同業 EV 含租賃；Microsoft EBITDA 已扣營業租金、未扣融資租賃折舊 → 減融資租賃負債")
        b.row(v, "sotp", "分部加總每股價值（目標價時點）", "US$", lambda y: f"({R('sev')}+{R('sbr')})/{R('sh_end', 2027)}", fmt="#,##0.00", scalar=True)
        b.row(v, "tp_s", "情境目標價＝w_dcf × DCF ＋（1 − w_dcf）× 分部加總", "US$", lambda y: f"w_dcf*{R('dcf12')}+(1-w_dcf)*{R('sotp')}", fmt="#,##0.00", scalar=True)
    return b


MAIN_ROUTE = {"thesis": "Thesis", "seg": "Seg", "pl": "PL", "cf": "CF", "val": "Val"}
SENS_ROUTE = {k: "Sens" for k in MAIN_ROUTE}
BL = {}


def build_block(s, title, route, P=None, V=None, base=None, ops=True, val=True):
    b = Block(s, title, route, P, V, base)
    engine(b, val=val, ops=ops)
    b.layout_and_write()                                  # 公式在寫入時才求值：此時全部列號已配置，可前後引用
    return b


for s in SC:
    BL[s] = build_block(s, SCN[s], MAIN_ROUTE)

# =====================================================================================
# 各頁標題與摘要（主模型）
# =====================================================================================
TH = WS["Thesis"]
TH["A1"] = "Thesis — 命題模組：AI 在役 MW → Azure AI 營收（MW × Tokenomics 每 MW 持有成本 × k）；AI 資本支出（MW × Tokenomics 每 MW 成本）→ AI 折舊與營運成本 → Agents and Infra 營業利益；AI 增量 ROIC 對 WACC"
TH["A2"] = "三情境的 MW 路徑相同（資本支出依公司指引），差在變現：k（長約／現貨組合）與 Copilot 座席。命題模組只驅動 Azure 的「Azure AI」子列與 M365 cloud 的「Copilot」子列，不取代官方分部"
TH["A3"] = "核心問題的答案列：「AI 增量 ROIC − WACC」（FY2029 欄）"
SG = WS["Seg"]
SG["A1"] = "Seg — 營收分部層：官方分部 Agents and Infra、Devices and Consumer（8-K 2026-09-02 重編）→ 八條產品線；Azure ＝ Azure AI（命題）＋ Azure 其他；Microsoft 365 cloud ＝ 基礎 ＋ Copilot（命題）"
SG["A2"] = "非命題產品線：FY2026 對 FY2025 年增線性收斂到長期值（FY2032 到達；In 頁 F 節）；三情境相同。FY2025、FY2026 欄為實績（Azure、M365 cloud 的子列為模型拆分）"
PLW = WS["PL"]
PLW["A1"] = "PL — 分部營收 → 分部營業利益（Agents and Infra 由核心利益率扣 AI 營運成本與 AI 折舊；Devices and Consumer 依實績利益率）→ 合併營業利益 → 業外 → 稅 → 淨利 → EPS"
PLW["A2"] = "核對列（三情境、全部年度合計，應為 0；segments.json 與 tools/segment_check 讀這三列）"
hdr(PLW, 3, YHDR)
CHKS = (("檢查：分部合計 − 合併營收（三情境，實績期另比對公司值）", lambda b, y: f"ABS({b.ref('p_ai_seg', y)}+{b.ref('p_dc', y)}-{b.ref('p_tot', y)})" + (f"+ABS({b.ref('p_tot', y)}-{qa('rev', y)})" if y <= 2026 else "")),
        ("檢查：產品線合計 − 分部（三情境，實績期另比對公司值）", lambda b, y: f"ABS({b.ref('c_seg', y)})"),
        ("檢查：命題子列合計 − 母列（Azure、Microsoft 365 cloud；三情境）", lambda b, y: f"ABS({b.ref('p_s_azai', y)}+{b.ref('p_azo', y)}-{b.ref('p_azure', y)})+ABS({b.ref('p_m365b', y)}+{b.ref('p_cop', y)}-{b.ref('p_m365', y)})"),
        ("檢查：分部營業利益合計 − 合併營業利益（三情境；實績期比對公司值）", lambda b, y: f"ABS({b.ref('ai_oi', y)}+{b.ref('dc_oi', y)}-{b.ref('oi', y)})" + (f"+ABS({b.ref('oi_chk', y)})" if y <= 2026 else "")))
PLCHK = []
for i, (lab, f) in enumerate(CHKS):
    r_ = 4 + i
    PLW.cell(r_, 1, lab); PLW.cell(r_, 2, "US$B")
    for y in range(2025, 2033):
        PLW[f"{col(y)}{r_}"] = "=" + "+".join(f(BL[s], y) for s in SC); PLW[f"{col(y)}{r_}"].number_format = "0.000"
    for cc in range(1, 12): PLW.cell(r_, cc).fill = CHK
    PLCHK.append(lab)
PLW["A9"] = "以下三個區塊依序為情境 A、B、C；分部與產品線列標籤與 8-K 相同（segments.json 對應）"

CFW = WS["CF"]
CFW["A1"] = "CF — 營業現金流（淨利＋AI 與非 AI 折舊＋股票薪酬＋營運資金）− 資本支出（AI＝Thesis；非 AI＝營收 × 比例）− 融資租賃本金 → 自由現金流（財報現金）；股利、回購、舉債（最低現金）"
CFW["A2"] = "收支關卡（地圖讀此表）：模型期間首年起五年（FY2027–FY2031）自由現金流扣股利後合計（財報現金，不含自用 AI 影子效益）；並判斷是否需要舉債或停止回購"
hdr(CFW, 4, ["項目", "情境 A（樂觀）", "情境 B（基準）", "情境 C（保守）", "說明"])
CFL = ["收支關卡：自由現金流扣股利（五年合計）", "是否需要舉債或停止回購（FY2027–FY2031）", "自由現金流 FY2027（參考）", "自由現金流 FY2029（參考）",
       "資本支出 FY2027（公司口徑，參考）", "回購合計 FY2027–FY2031", "新增舉債合計 FY2027–FY2031", "檢查：收支關卡＝首年起五年自由現金流扣股利（逐欄加總）"]
for i, lab in enumerate(CFL, 5):
    CFW.cell(i, 1, lab)
for j, s in enumerate(SC):
    c_ = "BCD"[j]; b = BL[s]
    CFW[f"{c_}5"] = "=" + b.ref("gate")
    CFW[f"{c_}6"] = "=" + b.ref("need")
    CFW[f"{c_}7"] = "=" + b.ref("fcf", 2027); CFW[f"{c_}8"] = "=" + b.ref("fcf", 2029); CFW[f"{c_}9"] = "=" + b.ref("capex", 2027)
    CFW[f"{c_}10"] = f"=SUM({b.rng('bb', 2027, 2031)})"; CFW[f"{c_}11"] = f"=SUM({b.rng('nd', 2027, 2031)})"
    CFW[f"{c_}12"] = f"={c_}5-(" + "+".join(b.ref("fcfd", yy) for yy in range(2027, 2032)) + ")"
    for rr in (5, 7, 8, 9, 10, 11):
        CFW[f"{c_}{rr}"].number_format = "#,##0.0"
    CFW[f"{c_}12"].number_format = "0.000"
CFW["E5"] = "US$B；為正＝五年自由現金流足以支付股利（地圖收支關卡讀 C 欄＝基準情境）。修正一：v0.3 初版的範圍公式被讀成 F 欄到 CF 欄，誤加了 FY2032"
for cc in range(1, 6): CFW.cell(5, cc).fill = KEY
for cc in range(1, 6): CFW.cell(12, cc).fill = CHK

# =====================================================================================
# Val：0. 目標價（0.1 DCF、0.2 分部加總〔Comparable〕、0.3 結論）；1. 以下為各情境評價明細
# =====================================================================================
VW = WS["Val"]
VW["A1"] = "Val — 目標價（約 12 個月後，2027-10-27）＝DCF 50%＋分部加總 SOTP 50%（Comparable 在分部層做）；情境權重 A 25%／B 50%／C 25%（皆 [判斷] 預設，Andy 可改，In 頁黃底）。主值＝讀法 2（自用 AI 價值中性），讀法 1 並列"
VW["A2"] = ("口徑（修正一）：DCF 的自由現金流扣股票薪酬（視為現金成本）、加自用 AI 影子效益（稅後），在帳融資租賃負債自權益橋扣除，推到目標價時點後扣期間股利；"
            "SOTP 的產品線 EBITDA 為股票薪酬加回口徑（與同業倍數同口徑），AI 基礎設施模組以自身 DCF（含資本支出與汰換）計價，同業倍數版列對照")
VR_ = {}
vr = [4]


def vset(key, label, fs, fmt, note="", bold=False):
    r_ = vr[0]; VW.cell(r_, 1, label); VR_[key] = r_
    if bold:
        VW.cell(r_, 1).font = BOLD
    for c_, f in zip("BCD", fs or []):
        if f is None:
            continue
        VW[f"{c_}{r_}"] = f; VW[f"{c_}{r_}"].number_format = fmt
    if note:
        VW[f"E{r_}"] = note
    vr[0] += 1
    return r_


def vhead(key, labels):
    VR_[key] = vr[0]; hdr(VW, vr[0], labels); vr[0] += 1


def vtitle(key, text):
    VR_[key] = vr[0]; VW.cell(vr[0], 1, text).font = BOLD; vr[0] += 1


sc3 = lambda key: [f"={BL[s].ref(key)}" for s in SC]
B3 = lambda key: f"$B${VR_[key]}"
vtitle("t0", "0. 目標價")
vtitle("t01", "0.1 DCF（評價用自由現金流折現；WACC：Blume β、約 10 年期 kd、負債含融資租賃；g 3%）")
vhead("h01", ["項目", "情境 A（樂觀）", "情境 B（基準）", "情境 C（保守）", "說明"])
vset("ev", "DCF 企業價值（2026-06-30）", sc3("ev"), "#,##0.0", "US$B")
vset("tvs", "永續價值占比", sc3("tvs"), "0.0%", "永續價值現值 ÷ EV")
vset("nc", "加：淨現金（FY2026 年底）", [f"={BL[s].ref('netcash', 2026)}" for s in SC], "#,##0.0")
vset("flb", "減：在帳融資租賃負債（FY2026 年底）", [f"={BL[s].ref('fl_bal', 2026)}" for s in SC], "#,##0.0", "修正一 F6：融資租賃列入淨負債")
vset("oai", "加：OpenAI 持股（比例 × 估值 ×（1 − 流動性折價））", ["=oai_v"] * 3, "#,##0.0", "Anthropic 比例未揭露，不估值")
vset("ps0", "DCF 每股價值（2026-06-30）", sc3("ps0"), "#,##0.00", "÷ FY2026 稀釋股數")
vset("dcf12", "DCF 情境價值推到 12 個月後（×（1＋WACC）^t − 期間股利）", sc3("dcf12"), "#,##0.00", "t＝評價日 → 2027-10-27；扣 FY2027 股利與 FY2028 第一季股利")
r_dcf12 = VR_["dcf12"]
vset("dcftp", "DCF 目標價（情境加權）", [f"=w_A*B{r_dcf12}+w_B*C{r_dcf12}+w_C*D{r_dcf12}"], "#,##0.00", "Σ 情境權重 × 情境價值")
vset("wacc", "WACC／永續成長率", ["=wacc", "=g_tv"], "0.00%", "WACC＝In 頁（CAPM，Blume β）；g＝In 頁")
vr[0] += 1
vtitle("t02", "0.2 分部加總（SOTP；Comparable 在分部層做：非 AI 產品線＝目標價時點 NTM EBITDA × 同業 NTM EV/EBITDA 中位數；AI 基礎設施＝模組 DCF）")
vhead("h02", ["項目", "情境 A（樂觀）", "情境 B（基準）", "情境 C（保守）", "倍數（x）"])
SOTPL = (("ne_aimod", "AI 基礎設施模組（Azure AI＋自用 AI；模組 DCF）", None), ("ne_azo", "Azure 其他（非 AI 雲端同業）", "m_cloud"), ("ne_m365", "Microsoft 365 cloud（企業軟體同業）", "m_sw"),
         ("ne_lic", "Productivity and server licensing（企業軟體）", "m_sw"), ("ne_ind", "Industry solutions（企業軟體）", "m_sw"), ("ne_fro", "Frontier and support services（企業軟體）", "m_sw"),
         ("ne_srch", "Search and advertising（數位廣告同業）", "m_ads"), ("ne_xbox", "XBOX（遊戲同業）", "m_game"), ("ne_win", "Windows OEM and devices（PC／裝置同業）", "m_pc"))
for k, lab, m in SOTPL:
    r_ = vset(k, f"EV：{lab}", sc3(k), "#,##0.0")
    if m:
        VW[f"E{r_}"] = f"={m}"; VW[f"E{r_}"].number_format = "0.00"
vset("sev", "分部加總 EV", sc3("sev"), "#,##0.0", "US$B")
vset("sbr", "加：淨現金（FY2027 年底）＋OpenAI 持股 − 在帳融資租賃負債", sc3("sbr"), "#,##0.0")
vset("sotp", "分部加總每股價值（目標價時點）", sc3("sotp"), "#,##0.00", "÷ FY2027 年底流通股數")
r_sotp = VR_["sotp"]
vset("comp", "Comparable 目標價（市場主流同業法）", [f"=w_A*B{r_sotp}+w_B*C{r_sotp}+w_C*D{r_sotp}"], "#,##0.00", "＝分部加總（SOTP）情境加權；同業倍數見 In 頁 B 節與 E 欄")
vset("aicmp", "對照：Azure AI 以同業倍數計（NTM EBITDA × 6×；v0.3 初版口徑）", sc3("ne_azai_cmp"), "#,##0.0", "F3：AI 腿不再用 EBITDA × 倍數（忽略資本密集度）；差額見報告")
vr[0] += 1
vtitle("t03", "0.3 結論")
vhead("h03", ["項目", "數值", "對照", "對照", "說明"])
vset("tp", "目標價（DCF 與 Comparable 加權）", [f"=w_dcf*{B3('dcftp')}+(1-w_dcf)*{B3('comp')}"], "#,##0.00", "w_dcf × DCF ＋（1 − w_dcf）× 分部加總；[判斷] 預設 50／50；主值＝讀法 2")
vset("dcf2", "DCF 目標價（情境加權）", [f"={B3('dcftp')}"], "#,##0.00")
vset("cons", "市場共識目標價", ["=cons_tp", "=cons_tp_lo", "=cons_tp_hi"], "#,##0.00", "S&P（經 StockAnalysis）57 位平均；C、D＝最低、最高（2026-10-26）")
vset("px", "現價（2026-10-26 收盤）", ["=px"], "#,##0.00", "StockAnalysis（2026-10-26 收盤）")
vset("up", "上檔空間：加權目標價", [f"={B3('tp')}/px-1"], "0.0%", "目標價 ÷ 現價 − 1")
vset("upc", "上檔空間：Comparable 目標價", [f"={B3('comp')}/px-1"], "0.0%")
UP = B3("up")
vset("grade", "上檔空間級距", [f'=IF({UP}>0.5,"大於 50%",IF({UP}>0.25,"25%–50%",IF({UP}>=0,"0%–25%",IF({UP}>=-0.25,"0%～−25%",IF({UP}>=-0.5,"−25%～−50%","小於 −50%")))))'], "@",
     "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
RECF = lambda x: f'IF({x}>=0.5,"強力買進",IF({x}>=0.25,"買進",IF({x}>-0.25,"中立",IF({x}>-0.5,"放空","強力放空"))))'
vset("rec", "投資建議", ["=" + RECF(UP)], "@", "五級固定規則：≥ +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、≤ −50% 強力放空")
vset("tps", "情境目標價（DCF 50%＋SOTP 50%）", sc3("tp_s"), "#,##0.00", "A／B／C")
r_tps = VR_["tps"]
vset("ups", "上檔空間：情境目標價", [f"={c_}{r_tps}/px-1" for c_ in "BCD"], "0.0%")
vset("recb", "投資建議（基準情境 B，參考）", [None, "=" + RECF(f"C{VR_['ups']}")], "@", "地圖與評等用加權目標價")
vset("tp_r1", "目標價（讀法 1：自用 AI 只計成本；敏感度）", None, "#,##0.00", "Sens 頁「讀法 1」案例（三情境全部重算）")
vset("up_r1", "上檔空間（讀法 1）", None, "0.0%")
vset("rec_r1", "投資建議（讀法 1）", None, "@")
vset("pess", "悲觀情境價（C 情境 DCF，12 個月後）", [f"=D{r_dcf12}"], "#,##0.00", "工作單：悲觀情境價＝C 情境 DCF")
vset("band", "價格帶（DCF 不如預期～順利）", [f"=D{r_dcf12}", f"=B{r_dcf12}"], "#,##0.00", "B＝C 情境 DCF、C＝A 情境 DCF（皆推到 12 個月後）")
vset("pe28", "現價 ÷ FY2028 EPS（模型）", [f"=px/{BL[s].ref('eps', 2028)}" for s in SC], "0.0", "A／B／C；評價口徑 EPS（讀法 2）")
vset("tpe28", "加權目標價 ÷ FY2028 EPS（基準）", [f"={B3('tp')}/{BL['B'].ref('eps', 2028)}"], "0.0")
vset("meth", "目標價方法", None, "@")
VW[f"B{VR_['meth']}"] = ("DCF：評價用自由現金流（營業現金流 − 股票薪酬 − 公司口徑資本支出＋自用 AI 影子效益稅後），WACC（CAPM：rf 5.00%（2026-10-26）、Blume β、ERP 5%；kd＝rf＋10 年期 AAA 利差；負債含融資租賃），"
                         "FY2027–FY2032 明確期（FY2029 後 MW 新增速度收斂），永續期 NOPAT ×（1 − g ÷ WACC）÷（WACC − g），g 3%；評價日 2026-06-30，以 WACC 推到 2027-10-27 並扣期間股利。"
                         "SOTP：AI 基礎設施模組 DCF；Azure 其他 13.82×（非 AI 雲端 8 家）、M365 cloud／授權／產業／Frontier 10.8×（企業軟體 5 家，不含 MSFT）、搜尋與廣告 9.11×（6 家）、"
                         "XBOX 12.63×（遊戲 3 家 TTM）、Windows 14.47×（PC 3 家 TTM），乘目標價時點 NTM EBITDA（股票薪酬加回）；加淨現金與 OpenAI 持股（25% × 852 ×（1 − 20%））、減在帳融資租賃負債。")
for k in ("dcftp", "comp", "tp"):
    VW[f"B{VR_[k]}"].fill = KEY
for cc in range(1, 6): VW.cell(VR_["rec"], cc).fill = KEY
vr[0] += 1
assert vr[0] < 68, vr[0]
VW.cell(vr[0], 1, "1. 各情境評價明細（引擎輸出；列 70 起）").font = BOLD
VW.column_dimensions["A"].width = 62
for c_ in "BCD": VW.column_dimensions[c_].width = 15
VW.column_dimensions["E"].width = 70

# =====================================================================================
# Sens：龍捲風（加權目標價變動）＋判斷預設替代；下方為各案例的完整引擎區塊
# =====================================================================================
SW = WS["Sens"]
SW["A1"] = "Sens — 一次動一個驅動，三情境全部重算後的加權目標價（DCF 50%＋SOTP 50%、情境 25／50／25；主值讀法 2）；下方列 80 起為各案例的完整引擎區塊（與主模型同一個函式）"
hdr(SW, 3, ["案例", "驅動", "加權目標價", "相對基準", "相對基準 %", "說明"])
TP_BASE = f"Val!$B${VR_['tp']}"
CASES = [
    ("讀法 1（自用 AI 只計成本）", "命題：自用 AI 讀法", dict(reading="1"), {}, "F1：自用 AI MW 不計影子效益（v0.3 初版口徑）"),
    ("k −15%", "命題：每 MW 價格 k", dict(kmult="0.85"), {}, "三情境 k 同乘 0.85"),
    ("k +15%", "命題：每 MW 價格 k", dict(kmult="1.15"), {}, "三情境 k 同乘 1.15"),
    ("對外比例 50%", "命題：對外比例", dict(ext="ext_sh_lo"), {}, "自有 AI MW＝對外 ÷ 50%（同樣對外 MW 要多建自用）"),
    ("對外比例 75%", "命題：對外比例", dict(ext="ext_sh_hi"), {}, "C23：最大不確定"),
    ("MW 速度：Epoch 可辨識站點", "命題：MW 速度", dict(pace="pace_epoch", cap="cap_epoch"), {}, "每年 +650、上限 4,386（v0.2.2 保守情境）"),
    ("MW 速度 +20%", "命題：MW 速度", dict(pace="pace*1.2"), {}, "FY2029 前每年 +3,033，之後同樣收斂"),
    ("AI IT 折舊 5 年", "命題：AI 折舊年限", dict(life="life_it_lo"), {}, "折舊與汰換提前"),
    ("AI IT 折舊 7 年", "命題：AI 折舊年限", dict(life="life_it_hi"), {}, "—"),
    ("Copilot 座席淨增 −50%", "命題：Copilot 座席", dict(cop_add="cop_add_{s}*0.5"), {}, "各情境每年淨增座席減半"),
    ("Copilot 座席淨增 +50%", "命題：Copilot 座席", dict(cop_add="cop_add_{s}*1.5"), {}, "—"),
    ("Copilot 月費 18 美元", "命題：Copilot 月費", dict(cop_px="cop_px_lo"), {}, "牌價 30、基準實收 25（預測期）"),
    ("Copilot 月費 30 美元", "命題：Copilot 月費", dict(cop_px="cop_px_hi"), {}, "—"),
    ("Azure 其他長期年增 5%", "業務：Azure 其他成長", dict(g_azo="glt_azo_lo"), {}, "基準 8%"),
    ("Azure 其他長期年增 12%", "業務：Azure 其他成長", dict(g_azo="glt_azo_hi"), {}, "—"),
    ("搜尋與廣告長期年增 4%", "業務：Search and advertising", dict(g_srch="glt_srch_lo"), {}, "基準 7%"),
    ("搜尋與廣告長期年增 10%", "業務：Search and advertising", dict(g_srch="glt_srch_hi"), {}, "—"),
    ("XBOX 長期年增 −2%", "業務：XBOX", dict(g_xbox="glt_xbox_lo"), {}, "基準 2%"),
    ("XBOX 長期年增 5%", "業務：XBOX", dict(g_xbox="glt_xbox_hi"), {}, "—"),
    ("分部利益率 −2pt", "業務：分部利益率", dict(dm="dm_seg_lo"), {}, "Agents and Infra 核心利益率與 Devices and Consumer 利益率同減"),
    ("分部利益率 +2pt", "業務：分部利益率", dict(dm="dm_seg_hi"), {}, "—"),
    ("WACC −1pt", "評價：WACC", {}, dict(wacc="(wacc-0.01)"), "只改折現（營運不變）"),
    ("WACC +1pt", "評價：WACC", {}, dict(wacc="(wacc+0.01)"), "—"),
    ("g 2%", "評價：永續成長", {}, dict(g="g_tv_lo"), "—"),
    ("g 4%", "評價：永續成長", {}, dict(g="g_tv_hi"), "—"),
    ("永續期新投資報酬 WACC＋5pt", "評價：永續期新投資報酬", {}, dict(ronic="ronic_prem_hi"), "F5：RONIC＝WACC 為判斷，列敏感度"),
    ("SOTP 倍數 −20%", "評價：同業倍數", {}, dict(mm="0.8"), "非 AI 產品線倍數同乘（AI 模組為 DCF，不受影響）"),
    ("SOTP 倍數 +20%", "評價：同業倍數", {}, dict(mm="1.2"), "—"),
]
ALT = [
    ("C 情境改用 Epoch 可辨識站點 MW 速度（650／年，上限 4,386）", {"C": (dict(pace="pace_epoch", cap="cap_epoch"), {})}, None),
    ("C 情境 k 改用長約低端（k_long_lo，v0.2.2 價格軸低）", {"C": (dict(kmult="(k_long_lo/k_C)"), {})}, None),
    ("永續期改用 FY2032 評價用自由現金流 ×（1＋g）", {s: ({}, dict(ronic="0", g="g_tv")) for s in SC}, "tvalt"),
    ("非 AI 各線改用 Microsoft 自身倍數 15.67×", {s: ({}, dict(m_sw="m_own", m_cloud="m_own", m_ads="m_own", m_game="m_own", m_pc="m_own")) for s in SC}, None),
    ("OpenAI 持股 27%（重組時比例）", {s: ({}, dict(oai="(oai_stake_hi*oai_val*(1-liq_disc))")) for s in SC}, None),
    ("OpenAI 持股 27%、無流動性折價", {s: ({}, dict(oai="(oai_stake_hi*oai_val)")) for s in SC}, None),
]
CUR["Sens"] = 80
srow = 4
CASE_ROWS = []
TPF = lambda blks: "=w_dcf*(" + "+".join(f"w_{s}*{blks[s].ref('dcf12')}" for s in SC) + ")+(1-w_dcf)*(" + "+".join(f"w_{s}*{blks[s].ref('sotp')}" for s in SC) + ")"
CASE_BLK = {}
for lab, drv, P, Vv, note in CASES:
    blks = {}
    for s in SC:
        if P:
            blks[s] = build_block(s, f"敏感度「{lab}」｜{SCN[s]}", SENS_ROUTE, P=P, V=Vv)
        else:
            blks[s] = build_block(s, f"敏感度「{lab}」｜{SCN[s]}", SENS_ROUTE, V=Vv, base=BL[s], ops=False)
    CASE_BLK[lab] = blks
    SW.cell(srow, 1, lab); SW.cell(srow, 2, drv)
    SW.cell(srow, 3, TPF(blks)).number_format = "#,##0.00"
    SW.cell(srow, 4, f"=C{srow}-{TP_BASE}").number_format = "+#,##0.00;-#,##0.00"
    SW.cell(srow, 5, f"=C{srow}/{TP_BASE}-1").number_format = "+0.0%;-0.0%"
    SW.cell(srow, 6, note)
    CASE_ROWS.append(srow); srow += 1
SENS_LAST = srow - 1
CR = {c_[0]: r_ for c_, r_ in zip(CASES, CASE_ROWS)}
srow += 1
SW.cell(srow, 1, "判斷預設的替代（報告第 5 節）：每列只改一項，其餘同基準").font = BOLD; srow += 1
hdr(SW, srow, ["替代", "影響範圍", "加權目標價", "相對基準", "相對基準 %", "說明"]); srow += 1
ALT_ROWS = []
for lab, mp, kind in ALT:
    blks = {}
    for s in SC:
        if s in mp:
            P, Vv = mp[s]
            if P:
                blks[s] = build_block(s, f"判斷替代「{lab}」｜{SCN[s]}", SENS_ROUTE, P=P, V=Vv)
            else:
                blks[s] = build_block(s, f"判斷替代「{lab}」｜{SCN[s]}", SENS_ROUTE, V=Vv, base=BL[s], ops=False)
                if kind == "tvalt":                       # 永續期自由現金流改為 FY2032 評價用自由現金流 ×（1＋g）
                    sh_, r_, _ = blks[s].R["tvfcf"]
                    WS[sh_].cell(r_, 1, "永續期自由現金流（替代：FY2032 評價用自由現金流 ×（1＋g））")
                    WS[sh_].cell(r_, 3, f"={BL[s].ref('fcf_val', 2032)}*(1+g_tv)")
        else:
            blks[s] = BL[s]
    SW.cell(srow, 1, lab); SW.cell(srow, 2, "、".join(mp.keys()) if len(mp) < 3 else "三情境")
    SW.cell(srow, 3, TPF(blks)).number_format = "#,##0.00"
    SW.cell(srow, 4, f"=C{srow}-{TP_BASE}").number_format = "+#,##0.00;-#,##0.00"
    SW.cell(srow, 5, f"=C{srow}/{TP_BASE}-1").number_format = "+0.0%;-0.0%"
    ALT_ROWS.append(srow); srow += 1
VD, VC = VR_["dcf12"], VR_["sotp"]
for lab, f in (("DCF／SOTP 權重改 30／70", f"=0.3*Val!{B3('dcftp')}+0.7*Val!{B3('comp')}"), ("DCF／SOTP 權重改 70／30", f"=0.7*Val!{B3('dcftp')}+0.3*Val!{B3('comp')}"),
               ("情境權重改 30／50／20", f"=w_dcf*(0.3*Val!$B${VD}+0.5*Val!$C${VD}+0.2*Val!$D${VD})+(1-w_dcf)*(0.3*Val!$B${VC}+0.5*Val!$C${VC}+0.2*Val!$D${VC})"),
               ("只用基準情境 B", f"=Val!$C${VR_['tps']}")):
    SW.cell(srow, 1, lab); SW.cell(srow, 2, "評價權重"); SW.cell(srow, 3, f).number_format = "#,##0.00"
    SW.cell(srow, 4, f"=C{srow}-{TP_BASE}").number_format = "+#,##0.00;-#,##0.00"
    SW.cell(srow, 5, f"=C{srow}/{TP_BASE}-1").number_format = "+0.0%;-0.0%"
    ALT_ROWS.append(srow); srow += 1
SW.cell(srow + 1, 1, "讀法：龍捲風以加權目標價的變動排序；命題驅動（自用 AI 讀法、k、對外比例、MW 速度、AI 折舊年限、Copilot）與業務驅動、評價參數並列。判斷預設的替代見上表").font = BOLD
SENS_NOTE = srow + 1
assert SENS_NOTE < 80, SENS_NOTE
SW.column_dimensions["A"].width = 52; SW.column_dimensions["B"].width = 24
for c_ in "CDE": SW.column_dimensions[c_].width = 14
SW.column_dimensions["F"].width = 50
# Val：讀法 1 目標價（Sens 案例）
R1 = CR["讀法 1（自用 AI 只計成本）"]
VW[f"B{VR_['tp_r1']}"] = f"=Sens!C{R1}"; VW[f"B{VR_['tp_r1']}"].number_format = "#,##0.00"
VW[f"B{VR_['up_r1']}"] = f"={B3('tp_r1')}/px-1"; VW[f"B{VR_['up_r1']}"].number_format = "0.0%"
VW[f"B{VR_['rec_r1']}"] = "=" + RECF(B3("up_r1"))


# =====================================================================================
# Recon：逐分部對照共識與指引；v0.3 vs v0.2.2
# =====================================================================================
RC = WS["Recon"]
RC["A1"] = "Recon — 模型（基準情境 B）對照市場共識與公司指引；差距超過 5% 者附原因。產品線層級的市場共識：未找到（免費資料源只提供合併層級）"
hdr(RC, 3, ["項目", "模型（B）", "共識／指引", "差距", "差距 %", "說明"])
B_ = BL["B"]
rc = [4]
RCF = {}


def rrow(label, m, c, note="", fmt="#,##0.0", pct=True):
    r_ = rc[0]
    RC.cell(r_, 1, label)
    if m is not None:
        RC.cell(r_, 2, m).number_format = fmt
    if c is not None:
        RC.cell(r_, 3, c).number_format = fmt
    if m is not None and c is not None:
        RC.cell(r_, 4, f"=B{r_}-C{r_}").number_format = fmt
        if pct:
            RC.cell(r_, 5, f"=B{r_}/C{r_}-1").number_format = "+0.0%;-0.0%"
    RC.cell(r_, 6, note); rc[0] += 1
    hf = {"0.0%": "p1", "0.00": "n2", "0.000": "n3"}.get(fmt)
    if hf:
        RCF[str(r_)] = {"B": hf, "C": hf, "D": hf}
    return r_


def rsec(t):
    RC.cell(rc[0], 1, t).font = BOLD; rc[0] += 1


rsec("1. 合併層級對共識（MarketScreener 2026-10-08；FY2027 營收與 EPS 另有 S&P 2026-10-26）")
rrow("營收 FY2027（對 S&P）", f"={B_.ref('tot', 2027)}", "=cons_rev27", "觀點：Azure AI 依 MW × 每 MW 收入、其他線依近一年年增收斂")
rrow("營收 FY2028", f"={B_.ref('tot', 2028)}", "=cons_rev28", "—")
rrow("營收 FY2029", f"={B_.ref('tot', 2029)}", "=cons_rev29", "—")
for y in (2027, 2028, 2029):
    rrow(f"營業利益 FY{y}", f"={B_.ref('oi', y)}", f"=cons_ebit{y % 100}", "觀點：AI 資本支出依 Tokenomics 每 MW 成本與 6 年折舊進損益，折舊上升快於共識（見報告第 4 節）" if y == 2027 else "同上")
rrow("EPS FY2027（財報口徑；對 S&P non-GAAP）", f"={B_.ref('eps_r', 2027)}", "=cons_eps27", "口徑：模型為 GAAP（不含持股損益），S&P 為 non-GAAP", fmt="0.00")
rrow("EPS FY2028（MarketScreener GAAP）", f"={B_.ref('eps_r', 2028)}", "=cons_eps28", "—", fmt="0.00")
rrow("EPS FY2029", f"={B_.ref('eps_r', 2029)}", "=cons_eps29", "—", fmt="0.00")
for y in (2027, 2028, 2029):
    rrow(f"資本支出 FY{y}（模型公司口徑 vs 共識推定現金口徑）", f"={B_.ref('capex', y)}", f"=cons_cx{y % 100}", "拆法：模型含融資租賃（FY2026 融資租賃約 28）；FY2027 取公司指引推估下限 204.2" if y == 2027 else "觀點：MW 路徑 × 每 MW 成本＋汰換")
rsec("2. 分部與產品線對公司指引（FY27Q1 指引 × 4 只作量級參考；季節性未調整）")
rrow("Agents and Infra FY2027 ÷ 4 vs FY27Q1 指引中點", f"={B_.ref('ai_seg', 2027)}/4", "=g_q1_ai", "全年平均季對第一季：模型含年內成長，偏高屬拆法")
rrow("Devices and Consumer FY2027 ÷ 4 vs FY27Q1 指引中點", f"={B_.ref('dc', 2027)}/4", "=g_q1_dc", "XBOX 季節性（第二季旺季）")
rrow("合併營收 FY2027 ÷ 4 vs FY27Q1 指引中點", f"={B_.ref('tot', 2027)}/4", "=g_q1_rev", "—")
rrow("Azure 年增 FY2027 vs FY27Q1 指引中點", f"={B_.ref('azure', 2027)}/{B_.ref('azure', 2026)}-1", "=g_q1_az", "公司只給第一季；全年未給", fmt="0.0%", pct=False)
rrow("營業利益率變動 FY2027（模型）vs 公司「下降不到 1 個百分點」", f"={B_.ref('oim', 2027)}-{B_.ref('oim', 2026)}", "=g_fy27_opm", "觀點：模型 AI 折舊 FY2027 增加較多（AI IT 6 年、首年半年慣例）", fmt="0.0%", pct=False)
rrow("資本支出 FY2027 ÷ 4 vs FY27Q1 指引（>50，公司口徑）", f"={B_.ref('capex', 2027)}/4", "=g_q1_cx", "—")
rsec("3. 命題驗證列（C3：只作驗證，不回頭校準）")
rrow("Azure AI＋Copilot 營收 FY2026（模型）vs 公司 AI 業務年化（FY26Q3 >37）", f"={B_.ref('azai', 2026)}+{B_.ref('cop', 2026)}", "=ai_runrate",
     "公司數字為 FY26Q3 年化（期末水準），模型為 FY2026 全年平均；要對上須期末對外 MW 約多 20%–30% 或 k 較高")
rrow("Azure 實績 FY2026 中 Azure AI 占比（模型）", f"={B_.ref('azai', 2026)}/{B_.ref('azure', 2026)}", None, "公司未揭露 Azure 中 AI 占比（不存在）", fmt="0.0%")
rrow("Copilot 營收 FY2026（模型估計）", f"={B_.ref('cop', 2026)}", None, "公司未揭露 Copilot 營收（不存在）；座席 30 百萬（FY26Q4）")
rsec("4. 產品線層級市場共識")
for lab in ("Azure", "Microsoft 365 cloud", "Productivity and server licensing", "Industry solutions", "Frontier and support services", "Search and advertising", "XBOX", "Windows OEM and devices"):
    rrow(f"{lab}：市場共識", f"={B_.ref({'Azure': 'azure', 'Microsoft 365 cloud': 'm365', 'Productivity and server licensing': 'lic', 'Industry solutions': 'ind', 'Frontier and support services': 'fro', 'Search and advertising': 'srch', 'XBOX': 'xbox', 'Windows OEM and devices': 'win'}[lab], 2027)}",
         None, "未找到（FY2027 新產品線口徑 2026-09-02 才公布；免費資料源無產品線共識）")
rsec("5. v0.3 vs v0.2.2：目標價差異拆解（加總＝總變動）")
r0 = rc[0]
rrow("v0.2.2 基準目標價（讀法 1；DCF 45%＋EV/EBITDA 55%）", "=old_tp", None, "v0.2.2 成品")
rrow("① 評價權重 45／55 → 50／50（用 v0.2.2 的兩條腿）", "=0.5*old_dcf+0.5*old_ev-old_tp", None, "方法")
rrow("② DCF 腿：v0.2.2 DCF → v0.3 基準情境 DCF（× 50%）", f"=0.5*(Val!$C${VR_['dcf12']}-old_dcf)", None, "營收結構（分部＋產品線）、AI 折舊進損益與稅、資本支出路徑、永續期口徑")
rrow("③ 倍數腿：v0.2.2 EV/EBITDA → v0.3 基準情境 SOTP（× 50%）", f"=0.5*(Val!$C${VR_['sotp']}-old_ev)", None, "產品線 EBITDA × 各同業倍數（取代 AI 6×／非 AI 加權倍數）")
rrow("④ 情境加權：基準情境 → A 25%／B 50%／C 25%", f"=Val!$B${VR_['tp']}-Val!$C${VR_['tps']}", None, "v0.2.2 以基準情境為點位")
rrow("v0.3 加權目標價（修正一）", f"=Val!$B${VR_['tp']}", None, "—")
rrow("檢查：v0.2.2 ＋ ①＋②＋③＋④ − v0.3（應為 0）", f"=B{r0}+B{r0 + 1}+B{r0 + 2}+B{r0 + 3}+B{r0 + 4}-B{r0 + 5}", None, "—", fmt="0.000")
RC_V22 = r0
for cc in range(1, 7): RC.cell(r0 + 6, cc).fill = CHK
RC.column_dimensions["A"].width = 64
for c_ in "BCDE": RC.column_dimensions[c_].width = 13
RC.column_dimensions["F"].width = 80

# =====================================================================================
# Sources
# =====================================================================================
SO = WS["Sources"]
hdr(SO, 1, ["編號", "文件", "日期", "網址", "標記"])
SRC = [
    ("Microsoft Form 10-K FY2026（期末 2026-06-30，經查核）", "2026-07-29", F["is.rev.fy2026"]["url"], "Verified"),
    ("Microsoft 8-K 附件 99.1：FY27 Segments and Investor Metrics（分部重編、Azure 金額、FY27Q1 調整後指引）", "2026-09-02", F["guid.fy27q1.rev"]["url"], "Verified／Interested-party（指引）"),
    ("Microsoft FY26 各季財報新聞稿與法說（IR 網站）", "2025-10-29～2026-07-29", "https://www.microsoft.com/en-us/investor/", "Verified／Interested-party"),
    ("Microsoft 事實總帳（v0.1a 起；每筆附來源、日期、標記）", "2026-10-08", "microsoft/data/microsoft_facts_20261008.json", "依各筆"),
    ("Tokenomics v21.4 快照（20261024_Tokenomics_v21.4.xlsx，master 3ac70ac）", "2026-10-25", "microsoft/data/tokenomics_snapshot_v21.4.json", "Derived"),
    ("k 證據（IREN 長約、Silicon Data H100 指數、Azure 牌價）", "2026-10-08", "microsoft/data/kevidence_microsoft_20261008.json", "Analogy／Verified"),
    ("市場共識（MarketScreener、S&P 經 StockAnalysis）", "2026-10-08", "microsoft/data/consensus_msft_20261008.json", "Interested-party"),
    ("StockAnalysis 統計頁（收盤 532.60）", "2026-10-26", SA + "statistics/", "Verified"),
    ("StockAnalysis 預測頁（共識目標價 603.43、FY2027 營收 391.6、EPS 19.76）", "2026-10-26", SA + "forecast/", "Interested-party"),
    ("非 AI 雲端同業 NTM EV/EBITDA（8 家）", "2026-10-07", "microsoft/data/peers_cloud_20261008.json", "Derived"),
    ("數位廣告同業 NTM EV/EBITDA（6 家）", "2026-10-07", "microsoft/data/peers_ads_20261009.json", "Derived"),
    ("MAG 命題與驅動對照表 r1（C1–C23）", "2026-10-09", "mag/docs/plan/20261008_MAG_命題與驅動對照表_r1.md", "—"),
    ("v0.2.2 成品（對照）", "2026-10-25", "microsoft/dist/_待刪除/20261025_Microsoft收支模型_v0_2_2.xlsx", "Derived"),
    ("個股模型原則（命題＋營收分部）", "2026-10-27", "docs/plan/個股模型原則_命題加營收分部.md", "—"),
    ("MAG v0.3 工作單", "2026-10-27", "mag/docs/workorders/20261027_MAG_v0.3_個股模型重建_分部架構.md", "—"),
]
for i, (d, dt, u, t) in enumerate(SRC, 2):
    SO.cell(i, 1, i - 1); SO.cell(i, 2, d); SO.cell(i, 3, dt); SO.cell(i, 4, u); SO.cell(i, 5, t)
SO.cell(len(SRC) + 3, 1, "未找到：產品線層級市場共識；Copilot 營收與 FY2024、FY2025 年底座席；M365 商用座席數（事實總帳未收）；Azure 中 AI 占比；MSFT 個別長債殖利率；遊戲與 PC 同業 NTM 倍數（改用 TTM）")
SO.column_dimensions["B"].width = 90; SO.column_dimensions["C"].width = 22; SO.column_dimensions["D"].width = 80; SO.column_dimensions["E"].width = 24

# =====================================================================================
# README（命題、核心問題、投資建議、三條重點觀察〔公式〕、驅動→推導對照表）
# =====================================================================================
RD = WS["README"]
RB = BL["B"]
RD["A1"] = "Microsoft（MSFT）個股模型 v0.3 — 命題＋官方營收分部 → 損益表（2026-10-27；資料截至 FY26Q4；FY27Q1 財報 2026-10-28 未納入）"
RD["A2"] = "工作表：README → In → Q → Thesis → Seg → PL → CF → Val → Recon → Sens → Sources。金額十億美元；財年 6 月結束（FY2027＝2026-07～2027-06）"
RD["A3"] = "命題：AI 能否同時在 Azure（算力）與 Microsoft 365（Copilot 應用）變現，撐住翻倍的資本支出。"
RD["A4"] = "核心問題：FY2029 的 AI 增量營業利益能否覆蓋 AI 資本支出的資金成本（AI 增量 ROIC 對 WACC），且自由現金流不需停止回購。"
V_ = lambda k: f"Val!$B${VR_[k]}"
RD["A5"] = (f'="① 命題答案（FY2029，基準情境 B）：讀法 2（主值；自用 AI 價值中性）AI 增量 ROIC "&TEXT({RB.ref("roic", 2029)},"0.0%")&"，讀法 1（自用 AI 只計成本）"&TEXT({RB.ref("roic_r1", 2029)},"0.0%")&"，WACC "&TEXT(wacc,"0.0%")'
            f'&"；對外 Azure AI 本身 "&TEXT({RB.ref("roic_x", 2029)},"0.0%")&"。讀法 1 下 AI 容量越大、目標價越低（自用算力只有成本）；讀法 2 下容量增加對目標價中性偏正（敏感度：MW 速度改 Epoch 可辨識站點 "&TEXT(Sens!D{CR["MW 速度：Epoch 可辨識站點"]},"+0.0;-0.0")&"、+20% "&TEXT(Sens!D{CR["MW 速度 +20%"]},"+0.0;-0.0")&"）。AI 折舊 FY2026 "&TEXT({RB.ref("ai_da", 2026)},"0.0")&" → FY2029 "&TEXT({RB.ref("ai_da", 2029)},"0.0")&" 十億美元，財報口徑營業利益率 FY2029 "&TEXT({RB.ref("oim", 2029)},"0.0%")&"（FY2026 "&TEXT({RB.ref("oim", 2026)},"0.0%")&"）。"')
RD["A6"] = (f'="② 收支關卡：基準情境首年起五年（FY2027–FY2031）自由現金流扣股利合計 "&TEXT(CF!C5,"#,##0.0")&" 十億美元（樂觀 "&TEXT(CF!B5,"#,##0.0")&"、保守 "&TEXT(CF!D5,"#,##0.0")&"）；"&CF!C6&"。"')
RD["A7"] = (f'="③ 評價：加權目標價 "&TEXT({V_("tp")},"#,##0")&" 美元（讀法 2），較現價 "&TEXT(px,"#,##0.00")&" "&TEXT({V_("up")},"+0.0%;-0.0%")&"；DCF "&TEXT({V_("dcftp")},"#,##0")&"、分部加總 "&TEXT({V_("comp")},"#,##0")'
            f'&"；讀法 1 目標價 "&TEXT({V_("tp_r1")},"#,##0")&"（"&{V_("rec_r1")}&"）；市場共識 "&TEXT(cons_tp,"#,##0")&"。WACC "&TEXT(wacc,"0.0%")&"（−1pt 時目標價 "&TEXT(Sens!C{CR["WACC −1pt"]},"#,##0")&"）；'
            f'非 AI 各線改用 Microsoft 自身 NTM EV/EBITDA "&TEXT(m_own,"0.0")&"× 時目標價 "&TEXT(Sens!C{ALT_ROWS[3]},"#,##0")&"。"')
RD["A8"] = f'="投資建議："&{V_("rec")}&"（讀法 2 主值；上檔空間 "&TEXT({V_("up")},"+0.0%;-0.0%")&"；讀法 1："&{V_("rec_r1")}&"；五級固定規則：≥ +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、≤ −50% 強力放空）"'
RD["A9"] = '="收支關卡（地圖）：CF 頁「收支關卡：自由現金流扣股利（五年合計）」基準 "&TEXT(CF!C5,"#,##0.0")&" 十億美元（財報現金，不含影子效益）"'
for r_ in range(3, 10):
    RD.cell(r_, 1).alignment = WRAP
RD["A11"] = "驅動 → 推導對照表（每頁輸入、推導量、因果方向）"; RD["A11"].font = BOLD
hdr(RD, 12, ["頁", "輸入（In 頁鍵）", "推導量", "因果方向", "基準值（FY2029，情境 B）"])
DRV = [
    ("Thesis", "pace、mw_ext_26、ext_sh、rent_ext_f", "在役對外 AI MW、自有 AI MW（含自用）", "MW 越多 → Azure AI 營收↑、AI 資本支出與折舊↑", f"={RB.ref('ext', 2029)}"),
    ("Thesis", "tk_hold_*、p_long、p_spot、ls_*（k）", "每 MW 年收入（US$M/MW-年）", "k 越高 → Azure AI 營收與 EBITDA↑（三情境主要差異）", f"={RB.ref('rev_mw', 2029)}"),
    ("Thesis", "tk_cxit_*、tk_cxfac、self_build、capex_floor27、life_it", "AI 資本支出、AI 折舊", "資本支出 → 折舊 → Agents and Infra 營業利益↓；→ 自由現金流↓", f"={RB.ref('ai_da', 2029)}"),
    ("Seg", "cop_s26、cop_add_*、cop_px", "Copilot 營收（M365 cloud 子列）", "座席 × 月費 → M365 cloud↑", f"={RB.ref('cop', 2029)}"),
    ("Seg", "glt_*、n_conv", "非命題產品線營收", "近一年年增收斂到長期值", f"={RB.ref('tot', 2029)}"),
    ("PL", "（FY2026 校準）、dm_seg", "Agents and Infra 核心利益率、分部營業利益", "核心利益率 ×（分部 − Azure AI）＋ Azure AI − AI 成本 − AI 折舊", f"={RB.ref('ai_oi', 2029)}"),
    ("PL", "internal_ai_reading、tax_f、oie_f", "自用 AI 影子效益、淨利、EPS（評價口徑）", "讀法 2：自用 MW × 持有成本（扣 Copilot）進評價口徑營業利益", f"={RB.ref('eps', 2029)}"),
    ("CF", "sbc_r、wc27、cx_*、fl_pay_*、bb_target、dps_*", "自由現金流、回購、舉債、股數", "自由現金流扣股利 → 回購（不足先減）→ 舉債", f"={RB.ref('fcf', 2029)}"),
    ("Val", "rf、beta、erp、kd、g_tv、m_*、al_*、w_*", "DCF、SOTP（AI 模組 DCF）、加權目標價", "WACC↑ → DCF↓；倍數↑ → SOTP↑", f"={V_('tp')}"),
]
for i, rowv in enumerate(DRV, 13):
    for j, v in enumerate(rowv, 1):
        c_ = RD.cell(i, j, v)
        if j == 5:
            c_.number_format = "#,##0.00"
RD_DRV = (13, 13 + len(DRV) - 1)
RD["A23"] = "結果摘要（三情境）"; RD["A23"].font = BOLD
hdr(RD, 24, ["項目", "情境 A（樂觀）", "情境 B（基準）", "情境 C（保守）", "說明"])
SUMR = [("營收 FY2029（US$B）", "tot", 2029, "#,##0.0"), ("Azure AI 營收 FY2029", "azai", 2029, "#,##0.0"), ("Copilot 營收 FY2029", "cop", 2029, "#,##0.0"),
        ("營業利益率 FY2029（財報口徑）", "oim", 2029, "0.0%"), ("EPS FY2029（財報口徑，US$）", "eps_r", 2029, "0.00"), ("EPS FY2029（評價口徑，US$）", "eps", 2029, "0.00"),
        ("自由現金流 FY2029（財報現金）", "fcf", 2029, "#,##0.0"), ("AI 增量 ROIC FY2029（讀法 2）", "roic", 2029, "0.0%"), ("AI 增量 ROIC FY2029（讀法 1）", "roic_r1", 2029, "0.0%")]
for i, (lab, k, y, fm) in enumerate(SUMR, 25):
    RD.cell(i, 1, lab)
    for j, s in enumerate(SC):
        RD.cell(i, 2 + j, f"={BL[s].ref(k, y)}").number_format = fm
RD.cell(35, 1, "情境目標價（US$；讀法 2）")
for j in range(3):
    RD.cell(35, 2 + j, f"=Val!{'BCD'[j]}{VR_['tps']}").number_format = "#,##0.00"
RD["A37"] = "與 v0.2.2 的差異：v0.2.2 沿用 CoreWeave 架構（對外 AI MW × 每 MW 收入＋其餘業務成長率線、融資瀑布）；v0.3 改為官方兩分部、八條產品線，命題只驅動 Azure AI 與 Copilot 兩條子列，AI 資本支出經折舊進分部營業利益；評價改 DCF 50%＋SOTP 50%、情境 25／50／25。拆解見 Recon 第 5 節。"
RD["A38"] = "已知限制：AI 在役 MW、對外比例、Copilot 營收與座席歷史、產品線利潤公司皆未揭露（Assumed，給區間）；FY27Q1 財報（2026-10-28）未納入，將以 v0.3.1 滾入；Tokenomics 維持 v21.4。"
RD["A40"] = "修正一（2026-10-27，獨立核對後）：收支關卡範圍公式誤加 FY2032（改為首年起五年並加核對列）；自用 AI 讀法 2 為主值；DCF 扣股票薪酬、減融資租賃負債、扣期間股利；SOTP 的 AI 腿改用模組 DCF；WACC 改 Blume β、10 年期 kd、負債含融資租賃；FY2029 後 MW 新增速度收斂；汰換時更新在役世代。前後對照見報告「修正一」節。"
RD["A39"] = "建置：python3 microsoft/build_microsoft.py raw.xlsx → soffice --headless --calc --convert-to xlsx → microsoft/dist/；HTML：python3 tools/company_html/build_company_html.py microsoft"
for r_ in (37, 38, 39, 40):
    RD.cell(r_, 1).alignment = WRAP
RD.column_dimensions["A"].width = 120
for c_ in "BCDE": RD.column_dimensions[c_].width = 24

for n in ("Thesis", "Seg", "PL", "CF"):
    ws = WS[n]
    ws.column_dimensions["A"].width = 64; ws.column_dimensions["B"].width = 12
    for y in YRS: ws.column_dimensions[col(y)].width = 11
    ws.column_dimensions["L"].width = 70
    ws.freeze_panes = "C4" if n != "PL" else "C4"
SW.column_dimensions["L"].width = 50
for c_ in "GHIJK": SW.column_dimensions[c_].width = 11
for c_ in "FGHIJK": VW.column_dimensions[c_].width = 11
VW.column_dimensions["L"].width = 60

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "build", "raw", "microsoft.xlsx")
os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
wb.save(out)

# =====================================================================================
# html.json（與 xlsx 同一份列號；文字不含手寫數字）
# =====================================================================================
def plr(key, s, y0=2025, y1=2032):
    sh, r_, _ = BL[s].R[key]
    return f"@{sh}!{col(y0)}{r_}:{col(y1)}{r_}"


def met(name, key, fmt, chart=False, y0=2025, y1=2032, ymin=None):
    m = {"name": name, "fmt": fmt, "rows": {s: plr(key, s, y0, y1) for s in SC}}
    if chart:
        m["chart"] = True
    if ymin:
        m["ymin"] = ymin
    return m


thB = BL["B"].R
yrow = lambda s, key: BL[s].R[key][1]
HJ = {
    "company": "Microsoft", "ticker": "MSFT", "thesis": "@README!A3",
    "question": "FY2029 的 AI 增量營業利益能否覆蓋 AI 資本支出的資金成本（AI 增量 ROIC 對 WACC），且自由現金流不需停止回購",
    "scenarios": [{"key": s, "name": SCN[s]} for s in SC],
    "sections": [
        {"id": "thesis", "title": "命題與核心問題", "blocks": [
            {"type": "list", "items": ["@README!A5", "@README!A6", "@README!A7"]},
            {"type": "kv", "items": [["投資建議", f"@Val!B{VR_['rec']}"], [f"@Val!A{VR_['up']}", f"@Val!B{VR_['up']}|sp1"], [f"@Val!A{VR_['tp']}", f"@Val!B{VR_['tp']}|n2"],
                                     [f"@Val!A{VR_['tp_r1']}", f"@Val!B{VR_['tp_r1']}|n2"], [f"@Val!A{VR_['rec_r1']}", f"@Val!B{VR_['rec_r1']}"], [f"@Val!A{VR_['px']}", f"@Val!B{VR_['px']}|n2"],
                                     ["@CF!A5", "@CF!C5|n1"], ["@CF!A6", "@CF!C6"]]},
            {"type": "note", "parts": ["@README!A40"]},
            {"type": "note", "parts": ["@Thesis!A2"]}]},
        {"id": "drivers", "title": "驅動因子 → 推導量", "intro": "@Thesis!A1", "blocks": [
            {"type": "table", "head": ["@README!A12", "@README!B12", "@README!C12", "@README!D12", "@README!E12"],
             "rowgen": {"sheet": "README", "rows": f"{RD_DRV[0]}-{RD_DRV[1]}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "", "", "", "n2"]}},
            {"type": "table", "caption": "主要輸入（In 頁；標記沿用 Excel）", "head": ["@In!B3", "@In!D3", "@In!E3", "@In!F3", "@In!C3", "@In!G3"],
             "rows": [[f"@In!B{IR[k]}", f"@In!D{IR[k]}|{f}", f"@In!E{IR[k]}|{f}", f"@In!F{IR[k]}|{f}", f"@In!C{IR[k]}", f"@In!G{IR[k]}|tag"]
                      for k, f in (("mw_ext_26", "n0"), ("pace", "n1"), ("ext_sh", "p0"), ("self_build", "p0"), ("k_B", "n3"), ("ls_A", "p0"), ("ls_C", "p0"),
                                   ("cop_s26", "n0"), ("cop_px", "n0"), ("cop_add_B", "n0"), ("capex_floor27", "n1"), ("life_it", "n0"), ("glt_azo", "p0"), ("internal_ai_reading", "n0"), ("beta", "n3"), ("wacc", "p2"), ("g_tv", "p0"))]},
            {"type": "table", "caption": f"@Thesis!A{yrow('B', 'yr') - 2}", "head": [f"@Thesis!{c_}{yrow('B', 'yr') - 1}" for c_ in "ABFGHIJK"],
             "rows": [[f"@Thesis!A{yrow('B', k)}", f"@Thesis!B{yrow('B', k)}"] + [f"@Thesis!{col(y)}{yrow('B', k)}|{f}" for y in FCY]
                      for k, f in (("pace", "n0"), ("ext", "n0"), ("own", "n0"), ("rev_mw", "n2"), ("azai", "n1"), ("ai_capex", "n1"), ("ai_da", "n1"), ("shadow", "n1"), ("ai_ebit", "n1"), ("nppe", "n1"))]
             + [[f"@Thesis!A{yrow('B', k)}", f"@Thesis!B{yrow('B', k)}"] + [f"@Thesis!{col(y)}{yrow('B', k)}|p1" for y in FCY] for k in ("roic", "roic_r1", "roic_x")]
             + [[f"@Thesis!A{yrow('B', 'spread')}", f"@Thesis!B{yrow('B', 'spread')}"] + [f"@Thesis!{col(y)}{yrow('B', 'spread')}|sp1" for y in FCY]]}]},
        {"id": "pl", "title": "三情境損益（分部 → 合併）", "intro": "@PL!A1", "blocks": [
            {"type": "pl", "years": f"@PL!{col(2025)}{thB['p_tot'][1] - 0}:{col(2032)}{thB['p_tot'][1]}", "metrics": []},
        ]},
        {"id": "valuation", "title": "估值", "intro": "@Val!A1", "blocks": [
            {"type": "note", "parts": ["@Val!A2"]},
            {"type": "table", "caption": f"@Val!A{VR_['t03']}", "head": [f"@Val!{c_}{VR_['h03']}" for c_ in "ABCDE"],
             "rowgen": {"sheet": "Val", "rows": f"{VR_['tp']}-{VR_['tpe28']}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "n2", "n2", "n2", ""],
                        "row_fmt": {str(VR_[k]): v for k, v in (("up", {"B": "sp1"}), ("upc", {"B": "sp1"}), ("grade", {"B": ""}), ("rec", {"B": ""}), ("ups", {"B": "sp1", "C": "sp1", "D": "sp1"}),
                                                              ("recb", {"C": ""}), ("up_r1", {"B": "sp1"}), ("rec_r1", {"B": ""}), ("pe28", {"B": "x1", "C": "x1", "D": "x1"}), ("tpe28", {"B": "x1"}))}}},
            {"type": "bars", "title": "目標價與現價", "fmt": "n0", "items": [
                {"label": f"@Val!A{VR_['tp']}", "value": f"@Val!B{VR_['tp']}", "series": 0}, {"label": f"@Val!A{VR_['dcftp']}", "value": f"@Val!B{VR_['dcftp']}", "series": 1},
                {"label": f"@Val!A{VR_['comp']}", "value": f"@Val!B{VR_['comp']}", "series": 2}, {"label": f"@Val!A{VR_['tp_r1']}", "value": f"@Val!B{VR_['tp_r1']}", "series": 1},
                {"label": f"@Val!A{VR_['pess']}", "value": f"@Val!B{VR_['pess']}", "series": 1},
                {"label": f"@Val!A{VR_['cons']}", "value": f"@Val!B{VR_['cons']}", "series": 2}], "ref": {"label": f"@Val!A{VR_['px']}", "value": f"@Val!B{VR_['px']}"}},
            {"type": "table", "caption": f"@Val!A{VR_['t01']}", "head": [f"@Val!{c_}{VR_['h01']}" for c_ in "ABCDE"],
             "rowgen": {"sheet": "Val", "rows": f"{VR_['ev']}-{VR_['wacc']}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "n1", "n1", "n1", ""],
                        "row_fmt": {str(VR_["tvs"]): {"B": "p1", "C": "p1", "D": "p1"}, str(VR_["ps0"]): {"B": "n2", "C": "n2", "D": "n2"}, str(VR_["dcf12"]): {"B": "n2", "C": "n2", "D": "n2"},
                                    str(VR_["dcftp"]): {"B": "n2"}, str(VR_["wacc"]): {"B": "p2", "C": "p1"}}}},
            {"type": "table", "caption": f"@Val!A{VR_['t02']}", "head": [f"@Val!{c_}{VR_['h02']}" for c_ in "ABCDE"],
             "rowgen": {"sheet": "Val", "rows": f"{VR_['ne_aimod']}-{VR_['aicmp']}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "n1", "n1", "n1", "x1"],
                        "row_fmt": {str(VR_["sotp"]): {"B": "n2", "C": "n2", "D": "n2"}, str(VR_["comp"]): {"B": "n2"}}}},
            {"type": "note", "parts": [f"@Val!A{VR_['meth']}", "：", f"@Val!B{VR_['meth']}"]},
            {"type": "table", "caption": "@CF!A2", "head": ["@CF!A4", "@CF!B4", "@CF!C4", "@CF!D4"],
             "rowgen": {"sheet": "CF", "rows": "5-12", "cols": ["A", "B", "C", "D"], "fmts": ["", "n1", "n1", "n1"], "row_fmt": {"6": {"B": "", "C": "", "D": ""}, "12": {"B": "n3", "C": "n3", "D": "n3"}}}}]},
        {"id": "recon", "title": "對帳：共識、指引與前一版", "intro": "@Recon!A1", "blocks": [
            {"type": "table", "head": ["@Recon!A3", "@Recon!B3", "@Recon!C3", "@Recon!D3", "@Recon!E3", "@Recon!F3"],
             "rowgen": {"sheet": "Recon", "rows": f"4-{rc[0] - 1}", "cols": ["A", "B", "C", "D", "E", "F"], "fmts": ["", "n1", "n1", "n1", "sp1", ""], "row_fmt": RCF}},
            {"type": "table", "caption": "官方分部對帳（Q 頁核對列，實績應為零）", "head": ["@Q!A3", "@Q!J3", "@Q!K3", "@Q!L3", "@Q!M3"],
             "rows": [[f"@Q!A{QR[k]}", f"@Q!J{QR[k]}|n3", f"@Q!K{QR[k]}|n3", f"@Q!L{QR[k]}|n3", f"@Q!M{QR[k]}|n3"] for k in ("ai", "azure", "m365", "lic", "ind", "fro", "dc", "srch", "xbox", "win", "rev", "chk_rev", "chk_oi", "chk_y")]}]},
        {"id": "sens", "title": "敏感度", "intro": "@Sens!A1", "blocks": [
            {"type": "tornado", "title": "加權目標價相對基準的變動（一次動一個驅動；三情境全部重算）", "sheet": "Sens", "rows": f"4-{SENS_LAST}", "label_col": "A", "value_col": "E", "fmt": "sp1"},
            {"type": "table", "collapse": True, "summary": "展開：全部案例的加權目標價", "head": ["@Sens!A3", "@Sens!B3", "@Sens!C3", "@Sens!D3", "@Sens!E3", "@Sens!F3"],
             "rowgen": {"sheet": "Sens", "rows": f"4-{SENS_LAST}", "cols": ["A", "B", "C", "D", "E", "F"], "fmts": ["", "", "n2", "n2", "sp1", ""]}},
            {"type": "table", "caption": f"@Sens!A{ALT_ROWS[0] - 2}", "head": [f"@Sens!A{ALT_ROWS[0] - 1}", f"@Sens!B{ALT_ROWS[0] - 1}", f"@Sens!C{ALT_ROWS[0] - 1}", f"@Sens!D{ALT_ROWS[0] - 1}", f"@Sens!E{ALT_ROWS[0] - 1}"],
             "rowgen": {"sheet": "Sens", "rows": f"{ALT_ROWS[0]}-{ALT_ROWS[-1]}", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "", "n2", "n2", "sp1"]}}]},
        {"id": "status", "title": "資料狀態", "blocks": [
            {"type": "kv", "items": [["模型版本", ["$version", "（", "$date", "）"]], ["來源 xlsx", "$xlsx"], ["來源 xlsx 的雜湊指紋（取代建置 commit）", "$sha"],
                                     ["Tokenomics 快照版本", f"@In!H{IR['tk_hold_h100']}"], ["資料截止", "@README!A1"]]},
            {"type": "h3", "text": "已知限制"},
            {"type": "note", "parts": ["@README!A38"]},
            {"type": "h3", "text": "「未找到」清單"},
            {"type": "scan", "needle": "未找到", "collapse": True, "summary": "展開：標為未找到的輸入與來源", "sheets": {"In": ["B", "H", "I"], "Recon": ["A", "F"], "Sources": ["A", "B"]}, "empty": "沒有標為未找到的項目。"}]},
    ],
}
PLM = [met(["合併營收（", "@PL!B" + str(thB["p_tot"][1]), "）"], "p_tot", "n1", chart=True),
       met("Agents and Infra", "p_ai_seg", "n1"), met("Devices and Consumer", "p_dc", "n1"),
       met("@PL!A" + str(thB["p_azure"][1]), "p_azure", "n1", chart=True), met("@PL!A" + str(thB["p_s_azai"][1]), "p_s_azai", "n1"),
       met("@PL!A" + str(thB["p_m365"][1]), "p_m365", "n1"), met("@PL!A" + str(thB["p_cop"][1]), "p_cop", "n1"),
       met("@PL!A" + str(thB["ai_oi"][1]), "ai_oi", "n1"), met("@PL!A" + str(thB["oim"][1]), "oim", "p1", chart=True, ymin="auto"),
       met("@PL!A" + str(thB["p_shadow"][1]), "p_shadow", "n1", y0=2027, y1=2032) if False else met("@PL!A" + str(thB["oi_val"][1]), "oi_val", "n1"),
       met("@PL!A" + str(thB["eps_r"][1]), "eps_r", "n2"), met("@PL!A" + str(thB["eps"][1]), "eps", "n2", chart=True)]
PLC = [met("@CF!A" + str(BL["B"].R["fcf"][1]), "fcf", "n1", chart=True, y0=2027, y1=2032), met("@CF!A" + str(BL["B"].R["capex"][1]), "capex", "n1", y0=2027, y1=2032),
       met("@CF!A" + str(BL["B"].R["fcfd"][1]), "fcfd", "n1", y0=2027, y1=2032), met("@CF!A" + str(BL["B"].R["bb"][1]), "bb", "n1", y0=2027, y1=2032),
       met("@Thesis!A" + str(BL["B"].R["roic"][1]), "roic", "p1", chart=True, y0=2027, y1=2032, ymin="auto")]
HJ["sections"][2]["blocks"][0]["metrics"] = PLM
PLH = BL["B"].R["p_azure"][1] - 1
HJ["sections"][2]["blocks"][0]["years"] = f"@PL!{col(2025)}{PLH}:{col(2032)}{PLH}"
CFH = BL["B"].R["c_ni"][1] - 1
HJ["sections"][2]["blocks"].insert(1, {"type": "pl", "years": f"@CF!{col(2027)}{CFH}:{col(2032)}{CFH}", "metrics": PLC})
HJ["sections"][2]["blocks"].append({"type": "table", "caption": "@PL!A2", "head": ["@PL!A3", "@PL!D3", "@PL!E3", "@PL!F3", "@PL!G3", "@PL!H3", "@PL!I3", "@PL!J3", "@PL!K3"],
                                    "rowgen": {"sheet": "PL", "rows": "4-7", "cols": ["A", "D", "E", "F", "G", "H", "I", "J", "K"], "fmts": ["", "n3", "n3", "n3", "n3", "n3", "n3", "n3", "n3"]}})
json.dump(HJ, open(os.path.join(HERE, "html.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("寫出", out, "與 microsoft/html.json")
