# Microsoft v0.3.1 成品核對：公式錯誤 0、Q／Seg／PL／CF／Val／Recon 頁「檢查：」列全為 0、C10 對帳落差（基準情境）≈ 0、固定列標籤存在、
# 變動拆解 JSON 的最後一步＝本檔加權目標價。用法：python3 microsoft/check_v03.py <重算後的 xlsx>；失敗時非零結束。
import json
import os
import sys

import openpyxl

X = sys.argv[1]
HERE = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.load_workbook(X, data_only=True)
fail = []
ERR = ("#REF", "#DIV", "#VALUE", "#NAME", "#N/A", "#NUM", "Err:")
n_err = 0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(ERR):
                n_err += 1
                if n_err <= 10: fail.append(f"公式錯誤 {ws.title}!{c.coordinate} {c.value}")
print(f"公式錯誤：{n_err} 格")


def nums(ws, r):
    return [c.value for c in ws[r][1:] if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


n_chk = 0
for sh in ("Q", "Seg", "PL", "CF", "Val", "Recon"):
    ws = wb[sh]
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if not (isinstance(v, str) and v.startswith("檢查：")):
            continue
        vs = nums(ws, r)
        n_chk += 1
        if any(abs(x) > 1e-6 for x in vs):
            fail.append(f"{sh}!A{r}「{v}」有非零值 {[round(x, 6) for x in vs if abs(x) > 1e-6][:5]}")
print(f"核對列：{n_chk} 列（Q、Seg、PL、CF、Val、Recon；含三情境與 Sens 區塊）")
T = wb["Thesis"]
gaps = []
for r in range(1, T.max_row + 1):
    if str(T.cell(r, 1).value or "").startswith("C10 對帳落差"):
        gaps.append(T.cell(r, 6).value)
print(f"C10 對帳落差（主模型 A／B／C）：{[round(g, 4) for g in gaps[:3]]}")
if len(gaps) < 3 or abs(gaps[1]) > 1e-3:
    fail.append(f"C10 基準情境對帳落差不為 0：{gaps[:3]}")
V = wb["Val"]
labs = {V.cell(r, 1).value: r for r in range(1, V.max_row + 1)}
need = ["目標價（DCF 與 Comparable 加權）", "上檔空間：加權目標價", "投資建議", "上檔空間級距", "DCF 目標價（情境加權）", "Comparable 目標價（市場主流同業法）",
        "價格帶（DCF 不如預期～順利）", "市場共識目標價", "DCF 情境價值推到 12 個月後（×（1＋WACC）^t − 期間股利）"]
for n in need:
    if n not in labs: fail.append(f"Val 缺列「{n}」")
if not any(str(k).startswith("現價（2026-") for k in labs):
    fail.append("Val 缺「現價（日期 收盤）」列")
cf = wb["CF"]
if "收支關卡：自由現金流扣股利（五年合計）" not in {cf.cell(r, 1).value for r in range(1, cf.max_row + 1)}:
    fail.append("CF 缺「收支關卡：自由現金流扣股利（五年合計）」")
tp = V.cell(labs["目標價（DCF 與 Comparable 加權）"], 2).value
up = V.cell(labs["上檔空間：加權目標價"], 2).value
rec = V.cell(labs["投資建議"], 2).value
AJ = os.path.join(HERE, "data", "attrib_v031_20261029.json")
if os.path.exists(AJ):
    last = json.load(open(AJ, encoding="utf-8"))["steps"][-1]["tp"]
    if abs(last - tp) > 1e-6:
        fail.append(f"變動拆解 JSON 已過期：最後一步 {last:.4f} ≠ 本檔 {tp:.4f}（重跑 SENS=1 bash microsoft/make_v03.sh）")
else:
    fail.append("缺 microsoft/data/attrib_v031_20261029.json")
print(f"加權目標價 {tp:.2f}；上檔空間 {up:+.1%}；投資建議 {rec}")
if fail:
    print("失敗：")
    for f in fail: print("  -", f)
    sys.exit(1)
print("全部通過")
