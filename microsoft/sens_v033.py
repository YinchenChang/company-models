# Microsoft v0.3.3 敏感度（整本重算部分）：In 頁有低／高、但 Sens 頁原有活公式案例沒有涵蓋的鍵（含 IF_HoldEcon 四世代同時 he_all），
# 一次改一個（build_microsoft.py --set），LibreOffice 整本重算後讀 Val 頁加權目標價，寫成 JSON 供 build_microsoft.py 寫入 Sens 頁（原則第 9 節）。
# 用法：python3 microsoft/sens_v033.py [輸出 JSON]（預設 microsoft/data/sens_v033_20261030.json；環境變數 JOBS＝平行數，預設 2）
import json, os, shutil, subprocess, sys, tempfile
from concurrent.futures import ThreadPoolExecutor
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.join(HERE, "build_microsoft.py")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data", "sens_v033_20261030.json")
WD = tempfile.mkdtemp(prefix="msft_sens_")


def run(name, sets):
    raw = os.path.join(WD, f"{name}.xlsx"); rm = os.path.join(WD, f"{name}.json")
    args = sum((["--set", f"{k}={v}"] for k, v in sets), [])
    subprocess.run([sys.executable, BUILD, raw, "--no-html"] + args, check=True, stdout=subprocess.DEVNULL, env=dict(os.environ, ROWMAP=rm))
    out = os.path.join(WD, "out"); os.makedirs(out, exist_ok=True)
    dst = os.path.join(out, f"{name}.xlsx")
    for _ in range(3):                                    # LibreOffice 偶爾沒有輸出檔 → 重試
        prof = tempfile.mkdtemp(prefix="lo_msft_")
        subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, raw],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=dict(os.environ, LC_ALL="C.UTF-8"), timeout=900)
        shutil.rmtree(prof, ignore_errors=True)
        if os.path.exists(dst): break
    R = json.load(open(rm))
    wb = openpyxl.load_workbook(dst, data_only=True)
    V = wb["Val"]
    res = {"tp": V.cell(R["tp"], 2).value, "up": V.cell(R["up"], 2).value, "rec": V.cell(R["rec"], 2).value}
    assert isinstance(res["tp"], (int, float)), (name, res)
    return res, R


try:
    base, R0 = run("base", [])
    jobs = {"base|base": []}
    for k, sd, sets in R0["RCASES"]:
        jobs[f"{k}|{sd}"] = sets
    res = {"base|base": {**base, "set": []}}
    todo = [(k, v) for k, v in jobs.items() if k != "base|base"]
    with ThreadPoolExecutor(max_workers=int(os.environ.get("JOBS", "2"))) as ex:
        futs = {k: ex.submit(run, f"c{i:02d}", sets) for i, (k, sets) in enumerate(todo)}
        for k, f in futs.items():
            o, _ = f.result()
            res[k] = {**o, "set": jobs[k]}
            print(k, jobs[k], round(o["tp"], 2), f"{o['tp'] / base['tp'] - 1:+.2%}", o["rec"], flush=True)
    json.dump(res, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("base", base, "→", OUT)
finally:
    shutil.rmtree(WD, ignore_errors=True)
