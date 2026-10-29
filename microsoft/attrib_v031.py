# Microsoft v0.3 → v0.3.1 加權目標價變動拆解＋C10 重解（MW 速度）。
# 每一步都以 build_microsoft.py 建置變體、LibreOffice headless 整本重算後讀 Val 頁（不是近似公式）。
# 用法：python3 microsoft/attrib_v031.py [--solve]   （--solve：先重解 C10 MW 速度並寫回 build_microsoft.py 的 PACE_V031／PACE_NV0）
#       輸出 microsoft/data/attrib_v031_20261029.json（build_microsoft.py 讀入 Recon 第 5、6 節）
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BUILD = os.path.join(HERE, "build_microsoft.py")
OUT = os.path.join(HERE, "data", "attrib_v031_20261029.json")
WD = tempfile.mkdtemp(prefix="msft_attr_")
V03_TP = 385.56245022499   # v0.3 成品（dist/_待刪除/20261027_Microsoft個股模型_v0.3.xlsx）Val!B36


def run(name, args):
    raw = os.path.join(WD, f"{name}.xlsx")
    env = dict(os.environ, ROWMAP=os.path.join(WD, f"{name}.json"))
    subprocess.run([sys.executable, BUILD, raw, "--no-html"] + args, check=True, stdout=subprocess.DEVNULL, env=env)
    out = os.path.join(WD, "out"); os.makedirs(out, exist_ok=True)
    prof = tempfile.mkdtemp(prefix="lo_msft_")
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, raw],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=dict(os.environ, LC_ALL="C.UTF-8"), timeout=600)
    shutil.rmtree(prof, ignore_errors=True)
    rm = json.load(open(os.path.join(WD, f"{name}.json")))
    wb = openpyxl.load_workbook(os.path.join(out, f"{name}.xlsx"), data_only=True)
    V, T, S, P, C = wb["Val"], wb["Thesis"], wb["Seg"], wb["PL"], wb["CF"]
    res = dict(tp=V.cell(rm["tp"], 2).value, tp_r1=V.cell(rm["tp_r1"], 2).value, up=V.cell(rm["up"], 2).value, rec=V.cell(rm["rec"], 2).value,
               gap=T.cell(rm["gap27"], 6).value, gapA=T.cell(rm["gapA"], 6).value, gapC=T.cell(rm["gapC"], 6).value)
    sg = rm["segB"]
    res["fy27"] = {"ai": S.cell(sg["ai_seg"], 6).value, "dc": S.cell(sg["dc"], 6).value, "rev": S.cell(sg["tot"], 6).value,
                   **{k: S.cell(sg[k], 6).value for k in ("azure", "m365", "lic", "ind", "fro", "srch", "xbox", "win")},
                   "ai_oi": P.cell(sg["ai_oi"], 6).value, "dc_oi": P.cell(sg["dc_oi"], 6).value, "oi": P.cell(sg["oi"], 6).value, "capex": C.cell(rm["cxB"], 6).value}
    print(f"  {name}: tp {res['tp']:.4f}  r1 {res['tp_r1']:.4f}  gap27 B {res['gap']:+.4f}（A {res['gapA']:+.4f}／C {res['gapC']:+.4f}）", flush=True)
    return res


def S(**kw):
    return sum((["--set", f"{k}={v}"] for k, v in kw.items()), [])


def solve(name, base_args, p0=2527.2, p1=2800.0):
    """C10：解 MW 速度使基準情境 FY2027 對帳落差＝0（割線法；落差對速度近似線性）"""
    g0 = run(f"{name}_p0", base_args + S(pace=p0))["gap"]
    g1 = run(f"{name}_p1", base_args + S(pace=p1))["gap"]
    for it in range(6):
        p2 = p1 - g1 * (p1 - p0) / (g1 - g0)
        g2 = run(f"{name}_it{it}", base_args + S(pace=round(p2, 4)))["gap"]
        p0, g0, p1, g1 = p1, g1, round(p2, 4), g2
        if abs(g2) < 5e-4:
            return p1
    raise SystemExit(f"C10 未收斂：{name}")


LEG = ["--legacy-q1", "--snap", "v21.4"] + S(px=532.60, oai_stake=0.25, oie_f=-0.4, pace=2527.2, nonnv=0) + ["--set", "roll_d2==DATE(2027,10,27)"]
src = open(BUILD, encoding="utf-8").read()
if "--solve" in sys.argv:
    print("C10 重解（nonnv＝0）")
    c_args = ["--snap", "v21.4"] + S(px=532.60, oai_stake=0.25, nonnv=0)
    p_nv0 = solve("c10_nv0", c_args)
    print("C10 重解（nonnv＝1，主值）")
    p_v031 = solve("c10_nv1", ["--snap", "v21.7"] + S(px=532.60, oai_stake=0.25, nonnv=1))
    src = re.sub(r"PACE_V031 = [\d.]+", f"PACE_V031 = {p_v031}", src)
    src = re.sub(r"PACE_NV0 = [\d.]+", f"PACE_NV0 = {p_nv0}", src)
    open(BUILD, "w", encoding="utf-8").write(src)
    print(f"寫回 PACE_V031 = {p_v031}、PACE_NV0 = {p_nv0}")
p_v031 = float(re.search(r"PACE_V031 = ([\d.]+)", src).group(1))
p_nv0 = float(re.search(r"PACE_NV0 = ([\d.]+)", src).group(1))

steps = []
print("變動拆解")
s0 = run("S0", LEG); steps.append(dict(key="0", **{k: s0[k] for k in ("tp", "tp_r1", "up", "rec")}))
if abs(s0["tp"] - V03_TP) > 1e-6:
    raise SystemExit(f"S0 未重現 v0.3：{s0['tp']} vs {V03_TP}")
A = ["--legacy-q1", "--snap", "v21.4"] + S(px=532.60, oai_stake=0.25, oie_f=-0.4, pace=2527.2, nonnv=0)
sa = run("Sa", A)
B = ["--snap", "v21.4"] + S(px=532.60, oai_stake=0.25, oie_f=-0.4, pace=2527.2, nonnv=0)
sb = run("Sb", B)
C = ["--snap", "v21.4"] + S(px=532.60, oai_stake=0.25, pace=p_nv0, nonnv=0)
sc = run("Sc", C)
Dd = ["--snap", "v21.7"] + S(px=532.60, oai_stake=0.25, pace=p_v031, nonnv=1)
sd = run("Sd", Dd)
E = S(oai_stake=0.25)
se = run("Se", E)
sf = run("Sf", [])
for k, r in (("a", sa), ("b", sb), ("c", sc), ("d", sd), ("e", se), ("f", sf)):
    steps.append(dict(key=k, **{x: r[x] for x in ("tp", "tp_r1", "up", "rec")}))
out = dict(asOf="2026-10-29", note="每步累加一組變更後整本重算；S0＝v0.3 重現（--legacy-q1、Tokenomics v21.4、v0.3 輸入）", v03_tp=V03_TP,
           pace=dict(v03=2527.2, nv0=p_nv0, v031=p_v031), gap_final=dict(B=sf["gap"], A=sf["gapA"], C=sf["gapC"]),
           steps=steps, v03_fy27=s0["fy27"], v031_fy27=sf["fy27"])
json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("寫出", OUT)
for st in steps:
    print(st)
shutil.rmtree(WD, ignore_errors=True)
