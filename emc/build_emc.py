# 台光電（2383）收支與估值模型 v0.2.1（AI 半導體 Project：載板與 PCB 次產業；高速 CCL 銅箔基板）
# 用法：python3 emc/build_emc.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；每 GW CCL 金額以美元輸入、以匯率 fx 換算；EPS：新台幣。
# 口徑：歸屬母公司淨利。樣板：ibiden/（次產業）；台股寫法（月營收、法說）比照 delta/、quanta/。
# 與 Ibiden 不同：材料端不以加速器顆數驅動，改以「部署 GW × 每 GW 高速 CCL 金額（加速器／運算板、交換器板、UBB／中板）× 份額」，
# 並加「產能上限」（全產全銷：營收＝MIN（需求面、月產能 × 每張營收））。
import sys, os, glob
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ITC = None
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    if r[0].value == "AI IT 資本支出（加速器 × 每顆 IT capex）":
        ITC = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert ITC and all(v is not None for v in ITC[0]), ITC
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("racks", "capit"):
        BR[k] = (r[4].value, r[7].value, f"{rel(BR_F)}｜In!E{r[0].row}（GB300）、H{r[0].row}（VR200）")
assert len(BR) == 2, BR.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B），每 GW CCL 金額為百萬美元；產能：百萬張／月；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝台光電法說（經富果、鉅亨等轉述）與月營收公告；公司的產品組合、份額與毛利率指引為 Interested-party；券商估計皆為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
FG0312 = "https://blog.fugle.tw/post/earnings-call-2383-2026-03-12"
FG0827 = "https://blog.fugle.tw/post/earnings-call-2383-2026-08-27"
UA0729 = "https://uanalyze.com.tw/articles/9219952430"
TN1005 = "https://finance.technews.tw/2026/10/05/emc-2383-202609-financial-report/"
KGI = "https://www.kgi.com.hk/zh-hk/-/media/files/kgishk/research-reports/tw-reports/emc-2383-tt20260729en.pdf"
SETN = "https://www.mirrormedia.mg/external/setn_1915172"
MS = "https://www.itiger.com/news/1142209073"
YD = "https://finance.sina.cn/2026-09-28/detail-initkarh2920455.d.html"

sec("A. AI 部署規模（讀 markets/network、markets/chip_bridge、tsmc 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, ITC[0][j], None, "Derived", ITC[1],
        "加速器顆數 × 每顆 IT capex（含 ASIC）；本模型以此除以每 GW IT 資本支出得到部署 GW（與台達電、廣達模型相同）" if y == 26 else "同上")
inp("itpg_gb", "每 GW IT 資本支出（GB300）", "US$B", None, BR["capit"][0], None, "Derived", BR["capit"][2], "Tokenomics IF_CapexIT（經橋接表；唯讀）")
inp("itpg_vr", "每 GW IT 資本支出（VR200；2027 起）", "US$B", None, BR["capit"][1], None, "Derived", BR["capit"][2], "同上；Rubin Ultra 與 ASIC 機架未另訂（Assumed：同 VR200）")
inp("vr26", "2026 年部署中 Rubin 世代占比（混合每 GW IT 資本支出與每 GW CCL 金額用）", "%", *SIL.vals("vr26"), "Interested-party", "TrendForce 2026 組合 Rubin 22%（經 markets/memory；與台達電模型相同）", "公司：M9 2026 下半年開始出貨、2027 年明顯放量（富果 2026-03-12）")
inp("racks_gb", "每 GW 機架數（GB300）", "架", None, BR["racks"][0], None, "Derived", BR["racks"][2], "Recon 用：每架 CCL 金額換算")
inp("racks_vr", "每 GW 機架數（VR200）", "架", None, BR["racks"][1], None, "Derived", BR["racks"][2], "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、台達電、Ibiden 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
inp("fx", "匯率（新台幣 ÷ 美元）", "NT$/US$", *SIL.vals("fx_twd"), "Derived", "與台達電、廣達模型相同（台達電 2Q26 法說換算約 32.7；2026 年初 31.4）", "10-23 收盤：未找到")

sec("B. 每 GW 高速 CCL 金額與台光電份額（本專案權威：每 GW CCL 金額、台光電份額，規則第 13 條）")
TOPDOWN = "由上而下（Assumed）：三項合計 2026 年約 US$350M／GW＝CCL 高階（M7 以上）市場約 US$8B ÷ 部署約 23 GW；市場量級取自 Prismark 2024 年 M7 以上 US$57 億、年增 40%（" + YD + "，源達證券轉引，Interested-party）"
inp("cmp_26", *SIL.row9("ccl_cmp_26")[1:])
inp("sw_26", *SIL.row9("ccl_sw_26")[1:])
inp("ubb_26", *SIL.row9("ccl_ubb_26")[1:])
IDX = (("cmpx", "加速器／運算板", ((1.20, 1.45, 1.80), (1.30, 1.60, 2.00), (1.30, 1.70, 2.20)), "Analogy",
        "每架 CCL 張數：GB200／GB300 約 164 張、VR200 約 578 張（源達證券）→ 每 GW 張數 VR200 ÷ GB300＝2.1 倍；Morgan Stanley 每架金額 → 每 GW 1.68 倍（Interested-party）",
        "2026 年已含 22% Rubin，所以 2027 ÷ 2026 小於世代倍數；ASIC（TPU V6 → V8 每台 144 → 200 張，+39%）"),
       ("swx", "交換器板", ((1.15, 1.40, 1.70), (1.25, 1.60, 2.00), (1.30, 1.75, 2.20)), "Assumed",
        "交換器板 40 層以上、M9＋Q 布（Rubin；源達證券）；1.6T 2027 年隨 M8／M9 放量（本土投顧 2026-10-01）", "網通模型：交換器市場 ÷ AI IT 資本支出 2026–27 持平（約 6.4%）；本列的上升來自材料等級"),
       ("ubbx", "UBB／中板／背板", ((1.30, 1.80, 2.50), (1.50, 2.30, 3.00), (1.50, 2.60, 3.50)), "Assumed",
        "VR200 中板、Rubin Ultra 78 層正交背板、Feynman 預計 144 層（源達證券）", "Morgan Stanley：M9 以上單價為基準 12.1 倍、M8 為 M7 兩倍以上（Interested-party）"))
for key, lab, vals, tag, s1, s2 in IDX:
    for y in (27, 28, 29):
        inp(f"{key}_{y}", *SIL.row9(f"ccl_{key}_{y}")[1:])
for y, v in ((27, (0.85, 0.95, 1.05)), (28, (0.75, 0.90, 1.00)), (29, (0.70, 0.85, 1.00))):
    inp(f"rel_{y}", f"台光電份額 20{y} ÷ 2026（相對值）", "x", *v, "Assumed",
        "Prismark：台光電高速材料市占 2025 年 38.6%（經 uanalyze 2026-07-29）、無鹵高速 2024 年 40.1%；公司自估 2025 年超過 50%（富果 2026-03-12），Interested-party" if y == 27 else "同上",
        "對手：台燿、斗山、生益、松下（M9）；台光電為 M8、M9 首家認證量產（公司，Interested-party）" if y == 27 else "—")
for y, v in ((27, (0.95, 1.00, 1.08)), (28, (0.88, 0.97, 1.05)), (29, (0.85, 0.93, 1.00))):
    inp(f"pp_{y}", f"CCL 價格指數 20{y}（2026＝1；原料轉嫁與短缺）", "x", *v, "Assumed",
        "公司：2026 Q2 全產品漲價至少 10%（銅箔、玻纖布上漲）；下半年不排除再漲（富果 2026-03-12、08-27）" if y == 27 else "同上",
        "高階 CCL 需求 2024–26 年複合成長 26%、供給 7%（本土投顧 2026-10-01）；HVLP4 銅箔 2028 年仍缺約 20%（Morgan Stanley）" if y == 27 else "新產能到位後價格回落（Assumed）")

sec("C. 產能（全產全銷：營收上限＝平均月產能 × 12 × 每張營收）")
inp("cap_26", "平均月產能 2026", "百萬張／月", 5.85, 6.0, 6.15, "Derived", "公司：2025 年 585 萬張／月；2026 Q1 中山 +30 萬、Q4 大園 +30 萬（富果 2026-03-12）；年底 615 萬（本土投顧 2026-10-01）",
    "凱基 2026-07-29：585 萬張／月（公司簡介）")
inp("cap_27", "平均月產能 2027", "百萬張／月", 7.0, 7.76, 8.2, "Derived", "公司：2027 年新增 330 萬張（Q1 120 萬、Q3 120 萬、Q4 90 萬；台灣、中國、馬來西亞），年底 945 萬（富果 2026-03-12）",
    "依開出月份加權：6.15＋1.2 × 10.5／12＋1.2 × 4.5／12＋0.9 × 1.5／12；崑山、中山各 60 萬已從 2026 Q4 延到 2027 Q1（設備交期），低值反映再延遲")
inp("cap_28", "平均月產能 2028", "百萬張／月", 9.5, 10.3, 10.8, "Derived", "公司：2028 年 1,110 萬張／月（富果 2026-08-27）；945 → 1,110 線性", "本土投顧 2026-10-01：2028 年底 1,110 萬")
inp("cap_29", "平均月產能 2029", "百萬張／月", 11.1, 11.8, 13.0, "Assumed", "公司未揭露 2029 年計畫：未找到", "低值＝2028 年底不再擴產")
for y, v in ((27, (1.10, 1.25, 1.45)), (28, (1.15, 1.40, 1.70)), (29, (1.20, 1.50, 1.85))):
    inp(f"mix_{y}", f"每張營收組合指數 20{y}（M8→M9、層數；不含價格；2026＝1）", "x", *v, "Assumed",
        "公司：M7 以上占營收 1Q26 55–60% → 2Q26 60–65%；M9 2026 下半年個位數占比、2027 年明顯放量（凱基 2026-07-29、富果 2026-03-12）" if y == 27 else "同上",
        "Morgan Stanley：M8 單價為 M7 兩倍以上、M9 以上為基準 12.1 倍（Interested-party）" if y == 27 else "—")

sec("D. 營收結構錨點（公司揭露比例 × 營收）")
inp("ai_25", "2025 年 AI 相關營收（AI 伺服器＋交換器）÷ 營收", "%", 0.45, 0.50, 0.55, "Assumed", "公司：2025 年基礎設施占 65–70%（凱基 2026-07-29）；4Q25 基礎設施中低軌衛星約 10%、其餘 AI 伺服器與交換器約 90%（含通用伺服器，富果 2026-03-12）",
    "讀法：基礎設施 67.5% × AI 伺服器＋交換器 75%（2Q26 拆分比例）；只用於 FY2025 欄拆分")
inp("ai_26", "2026 年 AI 相關營收（AI 伺服器＋交換器）÷ 營收", "%", 0.525, 0.585, 0.64, "Interested-party", "公司：2Q26 基礎設施 75–80%，其中 AI 伺服器 55–60%、通用伺服器 15–20%、交換器 15–20%、低軌衛星 5–10%（凱基 2026-07-29）",
    "公司：下半年基礎設施 80% 以上（富果 2026-08-27）；基準＝78% ×（57.5%＋17.5%）；交換器含非 AI 企業交換器（未拆）")
for y, v in ((27, (0.0, 0.08, 0.15)), (28, (0.0, 0.06, 0.12)), (29, (0.0, 0.05, 0.10))):
    inp(f"g_nai{y}", f"非 AI 營收年增 20{y}（通用伺服器、低軌衛星、手機、筆電、車用）", "%", *v, "Assumed",
        "2Q26 年增：手持 +42%、車用 +58%（凱基 2026-07-29，含漲價）；3Q26 指引：手持持平、車用個位數季減" if y == 27 else "同上",
        "通用伺服器、記憶體模組（DDR5 RDIMM）、低軌衛星持續成長（公司）；手機、筆電成熟" if y == 27 else "—")

inp("inf25", "2025 年基礎設施營收 ÷ 營收（4Q25 代理）", "%", 0.65, 0.675, 0.70, "Interested-party", "富果 4Q25 法說摘要（2026-03-12）：基礎設施 65–70%、行動裝置 20–25%、車用工業及其他 5–10%（公司）", "全年比例未揭露，以 4Q25 代理；基礎設施內低軌衛星約 10%、其餘為 AI 伺服器與交換器")
inp("mob25", "2025 年行動裝置營收 ÷ 營收（4Q25 代理）", "%", 0.20, 0.225, 0.25, "Interested-party", "同上", "—")
inp("inf26", "2026 年基礎設施營收 ÷ 營收", "%", 0.70, 0.75, 0.80, "Interested-party", "富果 2Q26 法說摘要（2026-08-27）：基礎設施占比自 3Q25 起逐季上升，2H26 達 80% 或更高（公司）", "全年＝4Q25 67.5% 與 2H26 80% 之間（Assumed）")
inp("mob26", "2026 年行動裝置營收 ÷ 營收", "%", 0.15, 0.18, 0.22, "Assumed", "4Q25 20–25%，基礎設施升到 80% 後行動裝置占比下降", "車用工業及其他＝差額")
sec("E. 毛利率")
inp("d_ai", "AI 相關（高速材料）毛利率高於其他營收的幅度", "pt", 0.04, 0.08, 0.12, "Assumed", "部門別毛利率未揭露：未找到；毛利率隨 M7 以上占比上升：2025 年 29.8% → 2Q26 33.9%（公司）",
    "2Q26 毛利率改善歸因高階產品組合與滿載（富果 2026-08-27）")

sec("F. FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設；3Q26 法說截至 2026-10-25 尚未舉行）")
inp("q3_rev", "3Q26 營收（月營收合計＝前 9 月 − 上半年）", "NT$B", None, "=141.227-80.342", None, "Verified", TN1005 + "：9 月 215.32 億、前 9 月 1,412.27 億（年增 103.7%）",
    "上半年 803.42 億（uanalyze 2026-07-29）；8 月 201.46 億（財報狗）")
inp("q4g", "4Q26 營收季增", "%", 0.0, 0.06, 0.12, "Assumed", "公司未給 4Q 指引（3Q 法說預計 10 月底，日期未找到）；9 月月營收 215.3 億 × 3＝646 億（季增 6%）", "4Q 大園 +30 萬張；崑山、中山延到 2027 Q1（公司）")
inp("q3_gm", "3Q26 毛利率", "%", 0.335, 0.35, 0.365, "Interested-party", "公司 3Q 指引 35% ± 1.5%（富果 2026-08-27）", "凱基 2026-07-29：3Q26 34.7%；本土投顧 35.08%")
inp("q4_gm", "4Q26 毛利率", "%", 0.335, 0.35, 0.37, "Assumed", "同 3Q 指引；下半年不排除再漲價（公司）", "—")
inp("q_opx", "3Q、4Q26 營業費用 ÷ 營收", "%", 0.06, 0.065, 0.07, "Assumed", "2Q26 6.9%（凱基：毛利 160.03 − 營業利益 127.34 億）、FY2025 9.6%（公司）", "營收成長快於費用")
inp("q_nop", "3Q、4Q26 業外收益（每季）", "NT$B", -0.5, 0.1, 0.5, "Assumed", "2Q26 0.05（含中國政府補助 1.35 億；凱基）", "1H26、FY2025 業外：未找到")
inp("q_nir", "3Q、4Q26 歸屬母公司淨利 ÷ 稅前", "%", 0.75, 0.77, 0.79, "Derived", "2Q26 98.73 ÷（127.34＋0.52）＝0.772（凱基）", "FY2025 淨利 ÷ 營業利益 146.49 ÷ 191.08＝0.767（stockanalysis）")
inp("q_sh", "股數（十億股；EPS 換算）", "十億股", None, "=15.213/42.46", None, "Derived", "1H26 淨利 152.13 億 ÷ EPS 42.46（uanalyze 2026-07-29）", "凱基：358.3 百萬股；現金股利 89.6 億 ÷ 25 元＝358 百萬股")

sec("G. 營業費用、業外、稅與股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.45, 0.6, 1.0, "Derived", "FY2026 錨點實測約 0.47（Q 頁 B26）", "FY2025：營業費用 +55% ÷ 營收 +46%＝1.2（擴產期，stockanalysis）")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.01, 0.0, 0.01, "Assumed", "無", "Sens 用 ±1 個百分點（台光電營業費用率約 7%）")
inp("nop_f", "業外收益（每年）", "NT$B", -1.0, 0.5, 2.0, "Assumed", "2Q26 0.05（凱基）", "—")
inp("nir_f", "歸屬母公司淨利 ÷ 稅前", "%", 0.75, 0.77, 0.79, "Derived", "同 q_nir", "—")
inp("sh_g", "股數年變動", "%", 0.0, 0.0, 0.005, "Assumed", "1H26 358.3 百萬股；FY2025 推導 351.5（淨利 ÷ EPS，加權）", "—")
inp("mf", "AI 部署係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 AI IT 資本支出約 2 倍", "±20%")

sec("H. 共識、研調與市場資料（只作 Recon 對帳，不作校準）")
inp("b_rev26k", "凱基 FY2026 營收（2026-07-29）", "NT$B", None, 191.123, None, "Interested-party", KGI, "—")
inp("b_eps26k", "凱基 FY2026 EPS（2026-07-29）", "NT$", None, 108.84, None, "Interested-party", KGI, "—")
inp("b_eps26c", "市場共識 FY2026 EPS（凱基報告所列，2026-07-29）", "NT$", None, 98.72, None, "Interested-party", KGI, "共識來源（Bloomberg 或其他）未載明")
inp("b_eps26l", "本土投顧 FY2026 EPS（2026-10-01）", "NT$", None, 120.77, None, "Interested-party", SETN + "（三立轉載，未具名本土投顧）", "3Q26 EPS 估 36.85")
inp("b_rev27k", "凱基 FY2027 營收", "NT$B", None, 306.405, None, "Interested-party", KGI, "—")
inp("b_rev27l", "本土投顧 FY2027 營收", "NT$B", None, 353.014, None, "Interested-party", SETN, "—")
inp("b_eps27k", "凱基 FY2027 EPS（2026-07-29）", "NT$", None, 190.64, None, "Interested-party", KGI, "目標價 7,625＝40 倍 2027 EPS")
inp("b_eps27c", "市場共識 FY2027 EPS（凱基報告所列，2026-07-29）", "NT$", None, 193.33, None, "Interested-party", KGI, "—")
inp("b_eps27l", "本土投顧 FY2027 EPS（2026-10-01）", "NT$", None, 204.85, None, "Interested-party", SETN, "目標價 6,145＝30 倍 2027 EPS")
inp("ms_rev28", "Morgan Stanley 隱含 FY2028 營收（2025–28 年複合成長 76%）", "NT$B", None, "=94.261*1.76^3", None, "Derived", MS + "（Deep News 轉述，Interested-party）", "以 FY2025 實績 942.61 億推算")
inp("ms_eps28", "Morgan Stanley 隱含 FY2028 EPS（目標價 8,400 ÷ 20 倍 2028 EPS）", "NT$", None, "=8400/20", None, "Derived", MS, "淨利 2025–28 年複合成長 117%（同上）")
inp("hs_sh25", "Prismark：台光電高速材料市占 2025", "%", None, 0.386, None, "Interested-party", UA0729 + "（公司轉述 Prismark）", "無鹵高速 2024 年 40.1%（富果 2026-03-12）")
inp("hs_sh25c", "公司自估：高速材料市占 2025", "%", None, 0.50, None, "Interested-party", FG0312 + "（「超過 50%」）", "利基市場（HDI、高速、AI 伺服器、交換器、低軌衛星）多數超過 50%（富果 2026-08-27）")
inp("ms_rk_gb", "Morgan Stanley：每架 CCL 金額 GB300", "US$", None, 6695, None, "Interested-party", MS, "VR200 18,750、VR300 35,601")
inp("ms_rk_vr", "Morgan Stanley：每架 CCL 金額 VR200", "US$", None, 18750, None, "Interested-party", MS, "—")
inp("sht_gb", "每架 CCL 張數 GB200／GB300", "張", None, 164, None, "Interested-party", YD + "（源達證券測算）", "Hopper 約 212 張")
inp("sht_vr", "每架 CCL 張數 VR200", "張", None, 578, None, "Interested-party", YD, "—")

sec("I. 價格與估值")
inp("px", "股價（2026-10-22 收盤）", "NT$", None, 5070, None, "Interested-party", "Yahoo 股市 2383（10-22 14:30 收盤 5,070，+1.6%）", "財報狗：10-22 收盤 5,070；10-23 收盤：未找到")
inp("pe_ttm", "近 12 月 P/E（10-22）", "x", None, 79.28, None, "Interested-party", "Yahoo 股市、財報狗（近四季 EPS 63.95）", "HiStock 10 月 79.28")

sec("J. 目標價情境權重（Val 第 0 節用；Andy 可改）")
for s_, w_, nm_ in (("A", 0.3, "順利"), ("B", 0.5, "平平"), ("C", 0.2, "不如預期")):
    inp(f"tp_w{s_}", f"目標價情境權重：情境 {s_}（{nm_}）", "權重", None, w_, None, "Assumed",
        "Andy 可改（2026-10-25 目標價規格預設 A／B／C＝0.3／0.5／0.2）", "三個權重合計應為 1；v0.2 起用於 DCF 與 Comparable 的情境加權（Val 第 0 節）")

# ---------------- v0.2 設定與新增輸入（附加在 In 既有輸入之後，不位移既有列） ----------------
V2 = {'name': '台光電',
 'unit': 'NT$B',
 'shunit': '十億股',
 'cur': 'NT$',
 'shg': 'sh_g',
 'epsfmt': '0.00',
 'pxfmt': '#,##0',
 'keys': {'rev': 'rev', 'ni': 'ni', 'sh': 'sh_', 'eps': 'eps'},
 'roe': ('=AVERAGE(0.2245,0.3092,0.3424)', 'stockanalysis 比率頁（Return on Equity）：FY2023 22.45%、FY2024 30.92%、FY2025 34.24%', '單一來源；近 12 月 48.38%（同站，2026-10-06）'),
 'cons_tp': (7803, 'stockanalysis forecast 頁（S&P Global）18 位分析師平均 7,803（中位數 7,998），頁面時間 2026-10-22', '第二來源：未找到；個股頁（2026-09-04 快取）7,517.81'),
 'cons_feps': ('=5415/32.17', 'stockanalysis 個股頁（快取 2026-09-04）：價格 5,415 ÷ Forward PE 32.17；較新資料：未找到'),
 'peers': [('台燿（Taiwan Union Technology）', '6274.TWO', 26.37, '2026-10-22'), ('聯茂（ITEQ）', '6213.TW', 27.49, '2026-10-22'), ('金像電（Gold Circuit）', '2368.TW', 19.57, '2026-10-22')],
 'peer_note': '高速 CCL 與 AI 伺服器 PCB 同業：台燿、聯茂（CCL）、金像電（下游 PCB）。生益科技（600183.SS）為 A 股，依規則不列；Doosan 電子材料、松下未單獨上市或為綜合企業，未列。非 A 股公司的同業名單不放中國 A 股（A 股有結構性溢價；Andy 2026-10-25 11:53）',
 'fwd_w': (0.16666666666666666, 0.8333333333333334),
 'fwd_note': '約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）',
 'pe10': None,
 'px_note': '現價＝In 頁 px（2026-10-22 收盤 5,070；10-23 收盤未找到）',
 'method': 'DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年平均約 29%（再投資率約 14%），FY2027–29 驅動因子、FY2030–32 收斂法＋永續；ROE 高使再投資率低、情境 A 維持 FY2029 高淨利率，DCF 偏高估。Comparable：台燿、聯茂、金像電前瞻 P/E '
           '中位數（只有三家，台光電歷來享溢價），偏低估。情境對應：A＝順利、B＝平平、C＝不如預期（原顆數情境）。',
 'readme_val': '估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），FY2027–FY2032 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，推到 12 個月後）＋0.5 × Comparable 目標價（台燿、聯茂、金像電 前瞻 P/E 中位數 × 12 個月後前瞻 EPS）；另列 '
               'Comparable 單獨目標價、市場共識目標價與差距拆解（EPS 因素、倍數因素）、DCF 價格帶、十年 P/E 分位檢查（只作檢查）。',
 'readme_add': ['v0.2（2026-10-25）：(1) PL 頁最下方新增延伸表，三情境延伸到 FY2032；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt；淨利率由 FY2029 收斂到 A＝FY2029、B＝FY2029 '
                '與十年平均的中點、C＝十年平均；十年平均淨利率未找到，依工作單以 FY2029 × 0.8 替代，Assumed；PL 延伸表灰底）。(2) Val 第 0 節目標價＝DCF 與 Comparable 加權（tp_w_dcf＝0.5）；v0.1.1 的「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留，不再是目標價。(3) Sens '
                '新增目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.4／0.8）。(4) Recon 新增目標價對帳。(5) 現價：10-23 收盤未找到（Yahoo 股市、鉅亨、HiStock、nstock 於 2026-10-25 仍為 10-22），依工作單沿用 2026-10-22 收盤 5,070。 '
                '新增內容一律放在各頁既有內容之後，FY2025–FY2029 既有數值不變。',
                '情境對應：A＝順利（顆數基準）、B＝平平（CY2028 顆數 ×0.8、CY2029 ×0.85）、C＝不如預期（CY2028–29 ×0.6），即原本的三個顆數情境。三情境在 FY2028 前完全相同（只差顆數係數，從 CY2028 起），DCF 三情境仍照算，差異來自 FY2028 之後（含 FY2030 起的收斂法）。',
                '先驗：本版未做先驗，無 Andy 合理 P/E 對照。']}
sec("J. 目標價（DCF 與 Comparable 加權）與收斂法參數（v0.2 新增；黃底＝Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（DCF 永續；情境 B 的 FY2032 營收年增收斂目標）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "Sens 頁 ±1pt（只動 DCF，PL 收斂目標不隨之變動）")
inp("tp_roe", "ROE（近三年平均）", "%", None, V2["roe"][0], None, "Derived", V2["roe"][1], V2["roe"][2])
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 改為 0.5：DCF／Comparable 各半）", "Sens 頁 0.4／0.8")
inp("fade_gA", "情境 A（順利）：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂，FY2032 達到")
inp("fade_gB", "情境 B（平平）：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C（不如預期）：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("nim10_f", "十年平均淨利率的替代：最後一年（FY2029）淨利率 ×（本值）", "x", None, 0.8, None, "Assumed", "十年（2016–2025）平均淨利率：未找到（免費資料商只到近五年）；依工作單以最後一年 × 0.8 替代", "淨利率收斂目標：A＝FY2029、B＝FY2029 與替代值的中點、C＝替代值（PL 延伸表）")
inp("cons_tp", "市場共識目標價（賣方平均）", V2["cur"], None, V2["cons_tp"][0], None, "Interested-party", V2["cons_tp"][1], V2["cons_tp"][2])
inp("cons_feps", "共識前瞻 12 月 EPS（＝價格 ÷ 前瞻 P/E）", V2["cur"], None, V2["cons_feps"][0], None, "Interested-party", V2["cons_feps"][1], "賣方共識，只作對帳與差距拆解，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣；股數為十億股）。藍字＝輸入；黑字＝公式。實績取自台光電法說（經轉述）、券商整理的季報數字與月營收公告"
Q["A2"] = "FY2026 錨點＝1H26 實績＋3Q26 營收實績（月營收）＋3Q、4Q 假設（毛利率、營業費用率、業外、淨利 ÷ 稅前；4Q 營收＝3Q × (1＋季增)）。1H26 營業利益＝2Q26 127.34 億＋1Q26（2Q26 ÷ 1.7863；玩股網單季營業利益季增 78.63%）。1H26 與 FY2025 的業外、稅前利益：未找到。3Q26 財報與法說（預計 10 月底）公布後更新"
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（營收實績）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1", "來源2／說明"])
S1 = "FY2025：stockanalysis 年報數字、uanalyze 2026-03（EPS 41.67、營收 942.6 億、毛利率 29.8%）；1H26：" + UA0729 + "、" + KGI
S2 = "3Q26 營收：月營收公告（TechNews 2026-10-05）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "營收", {"B": 94.261, "C": 80.342, "D": "=q3_rev", "E": "=D4*(1+q4g)"}, tag="Verified／Assumed（4Q 季增）", s2="FY2025 富果 943 億（同值）；1H26 富果：2Q26 472.7 億；" + S2)
qrow(5, "毛利", {"B": 28.120, "C": "=C4*0.3202", "D": "=D4*q3_gm", "E": "=E4*q4_gm"}, tag="Verified／Interested-party（3Q 指引）", s2="1H26 毛利率 32.02%（uanalyze）；2Q26 毛利 160.03 億（凱基）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2025 29.8%（公司）；2Q26 33.9%")
qrow(7, "營業費用", {"B": "=B5-B8", "C": "=C5-C8", "D": "=D4*q_opx", "E": "=E4*q_opx"}, tag="Derived（毛利 − 營業利益）／Assumed")
qrow(8, "營業利益", {"B": 19.108, "C": "=12.734+12.734/1.7863", "D": "=D5-D7", "E": "=E5-E7"}, tag="Verified／Derived", s2="FY2025 營益率 20.3%（富果）；2Q26 127.34 億（凱基）；1Q26 由玩股網季增率反推")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="2Q26 26.9%（凱基）")
qrow(10, "業外收益", {"D": "=q_nop", "E": "=q_nop"}, ann="", tag="Assumed", s1="—", s2="FY2025、1H26：未找到")
qrow(11, "稅前利益", {"D": "=D8+D10", "E": "=E8+E10"}, ann="", tag="Derived", s1="—", s2="FY2025、1H26：未找到")
qrow(12, "歸屬母公司淨利", {"B": 14.649, "C": 15.213, "D": "=D11*q_nir", "E": "=E11*q_nir"}, tag="Verified／Derived（3Q、4Q）", s2="FY2025 uanalyze 146.5 億；1H26 2Q26 98.73 億（凱基）")
qrow(13, "淨利 ÷ 營業利益", {c: f"={c}12/{c}8" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(14, "股數（十億股）", {"B": "=14.649/41.67", "C": "=q_sh", "D": "=q_sh", "E": "=q_sh"}, ann="=q_sh", tag="Derived", s2="淨利 ÷ EPS（FY2025 為加權平均）")
qrow(15, "EPS", {c: f"={c}12/{c}14" for c in "BCDE"}, fmt="0.00", ann="=F12/F14", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(16, "  公司揭露 EPS", {"B": 41.67, "C": 42.46}, fmt="0.00", ann="", tag="Verified", s2="1Q26 14.90、2Q26 27.55（udn、uanalyze）")
qrow(17, "AI 相關營收（＝營收 × AI 相關占比）", {"B": "=B4*ai_25"}, ann="=F4*ai_26", tag="Derived／Interested-party", s1="公司產品組合（In 頁 D 節）", s2="各季 AI 營收金額未揭露：未找到")
qrow(18, "非 AI 營收（＝營收 − AI 相關）", {"B": "=B4-B17"}, ann="=F4-F17", tag="Derived", s1="—", s2="通用伺服器、低軌衛星、手持（手機、筆電、記憶體模組）、車用與工業")
Q.cell(23, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(24, 1, "FY2026 營收年增"); Q["B24"] = "=F4/B4-1"
Q.cell(25, 1, "FY2026 營業費用年增"); Q["B25"] = "=F7/B7-1"
Q.cell(26, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B26"] = "=B25/B24"
Q.cell(27, 1, "FY2026 營收 ÷ 凱基（07-29）− 1"); Q["B27"] = "=F4/b_rev26k-1"
Q.cell(28, 1, "FY2026 EPS ÷ 本土投顧（10-01）− 1"); Q["B28"] = "=F15/b_eps26l-1"
Q.cell(29, 1, "3Q26 營收季增（實績；公司指引 15–20%）"); Q["B29"] = "=D4/(C4-33.07)-1"
for rr in range(24, 30): Q[f"B{rr}"].number_format = "0.0%" if rr != 26 else "0.00"
Q.cell(29, 3, "2Q26 營收 472.75 億＝上半年 803.42 − 1Q26 330.7")
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 12), ("sh", 14), ("eps", 15), ("ai", 17), ("nai", 18)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$F${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 90; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", cmpx="cmpx", swx="swx", ubbx="ubbx", rel="rel", pp="pp", cap="cap", mix="mix", mf="1", gn="g_nai", d="d_ai", opxa="0")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 部署與每 GW 高速 CCL 金額（需求面；D 欄＝FY2026 校準年）")
    row("itc", "AI IT 資本支出 × 情境顆數係數", "US$B", {c_(y): f"=itc_{y % 100}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("itpg", "每 GW IT 資本支出", "US$B", {"D": "=itpg_gb*(1-vr26)+itpg_vr*vr26", **{c_(y): "=itpg_vr" for y in BU[1:]}})
    row("gw", "部署 GW（IT 電力）", "GW", {c_(y): f"={c_(y)}{R['itc']}/{c_(y)}{R['itpg']}" for y in BU}, fmt="0.0")
    for k, lab, base, ix in (("vcmp", "每 GW CCL：加速器／運算板", "cmp_26", P["cmpx"]), ("vsw", "每 GW CCL：交換器板", "sw_26", P["swx"]), ("vubb", "每 GW CCL：UBB／中板／背板", "ubb_26", P["ubbx"])):
        row(k, lab, "US$M/GW", {"D": f"={base}", **{c_(y): f"={base}*{ix}_{y % 100}" for y in BU[1:]}}, fmt="#,##0")
    row("pgw", "每 GW 高速 CCL 金額（三項合計）", "US$M/GW", {c_(y): f"={c_(y)}{R['vcmp']}+{c_(y)}{R['vsw']}+{c_(y)}{R['vubb']}" for y in BU}, fmt="#,##0")
    row("mku", "AI 高速 CCL 市場（美元；不含價格指數）", "US$B", {c_(y): f"={c_(y)}{R['gw']}*{c_(y)}{R['pgw']}/1000" for y in BU})
    row("mkn", "AI 高速 CCL 市場（新台幣）", "NT$B", {c_(y): f"={c_(y)}{R['mku']}*fx" for y in BU}, fmt="#,##0.0")
    row("pp", "CCL 價格指數（FY2026＝1）", "x", {"D": "=1", **{c_(y): f"={P['pp']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh0", "台光電有效份額（FY2026 校準）＝ AI 相關營收 ÷ 市場", "%", {"D": f"=a26_ai/D{R['mkn']}"}, fmt="0.0%",
        note="Derived：分子為公司 AI 相關占比 × 營收；分母為本模型的 AI 高速 CCL 市場（含 ASIC）")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "台光電有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("aid", "AI 相關營收（需求面）＝ 市場 × 份額 × 價格指數", "NT$B", {c_(y): f"={c_(y)}{R['mkn']}*{c_(y)}{R['sh']}*{c_(y)}{R['pp']}" for y in BU}, fmt="#,##0.0")
    sec_("B. 產能上限（全產全銷）與營收")
    row("cap", "平均月產能", "百萬張／月", {"D": "=cap_26", **{c_(y): f"={P['cap']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("rps0", "每張營收（FY2026 校準）＝ 錨點營收 ÷（月產能 × 12）", "NT$／張", {"D": f"=a26_rev/(D{R['cap']}*12)*1000"}, fmt="#,##0",
        note="Derived：全產全銷，FY2026 營收全部由產能吸收；含 PP（膠片）與所有等級")
    row("mix", "每張營收組合指數（M8→M9、層數；FY2026＝1）", "x", {"D": "=1", **{c_(y): f"={P['mix']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("rcap", "產能上限營收 ＝ 月產能 × 12 × 每張營收 × 組合指數 × 價格指數", "NT$B",
        {c_(y): f"={c_(y)}{R['cap']}*12*$D${R['rps0']}*{c_(y)}{R['mix']}*{c_(y)}{R['pp']}/1000" for y in BU}, fmt="#,##0.0")
    row("nai", "非 AI 營收", "NT$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']}{y % 100})" for y in BU[1:]}}, fmt="#,##0.0")
    row("ai", "AI 相關營收 ＝ MIN（需求面、產能上限 − 非 AI）", "NT$B", {"C": "=a25_ai", "D": f"=D{R['aid']}",
        **{c_(y): f"=MIN({c_(y)}{R['aid']},{c_(y)}{R['rcap']}-{c_(y)}{R['nai']})" for y in BU[1:]}}, fmt="#,##0.0", note="產能不足時，非 AI 照成長率出貨、AI 取剩餘產能（Assumed）")
    row("bind", "產能是否受限", "", {c_(y): f'=IF({c_(y)}{R["aid"]}>{c_(y)}{R["rcap"]}-{c_(y)}{R["nai"]}+0.0001,"受限","未受限")' for y in BU[1:]}, fmt="General")
    row("rev", "營收", "NT$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.0")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("ninf", "  基礎設施：非 AI（一般伺服器、低軌衛星等）", "NT$B", {"C": f"=inf25*C{R['rev']}-C{R['ai']}", "D": f"=inf26*D{R['rev']}-D{R['ai']}", **{c_(y): f"={c_(y)}{R['nai']}*$D{r}/$D{R['nai']}" for y in BU[1:]}}, fmt="#,##0.0", note="基礎設施占比 × 營收 − AI；FY2027 起占非 AI 比例固定")
    row("infra", "基礎設施營收（AI＋非 AI）", "NT$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['ninf']}" for y in FYS}, fmt="#,##0.0", note="公司應用別：基礎設施（伺服器、交換器、低軌衛星）")
    row("mob", "行動裝置營收", "NT$B", {"C": f"=mob25*C{R['rev']}", "D": f"=mob26*D{R['rev']}", **{c_(y): f"={c_(y)}{R['nai']}*$D{r}/$D{R['nai']}" for y in BU[1:]}}, fmt="#,##0.0", note="公司應用別：行動裝置")
    row("aut", "車用、工業及其他營收（差額）", "NT$B", {c_(y): f"={c_(y)}{R['rev']}-{c_(y)}{R['infra']}-{c_(y)}{R['mob']}" for y in FYS}, fmt="#,##0.0", note="公司應用別：車用、工業及其他")
    row("segchk", "檢查：分部合計 − 合併營收", "NT$B", {c_(y): f"={c_(y)}{R['infra']}+{c_(y)}{R['mob']}+{c_(y)}{R['aut']}-{c_(y)}{R['rev']}" for y in FYS}, fmt="0.000")
    row("aish", "AI 相關 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("util", "需求面 ÷ 產能上限（AI 需求面＋非 AI）", "%", {c_(y): f"=({c_(y)}{R['aid']}+{c_(y)}{R['nai']})/{c_(y)}{R['rcap']}" for y in BU}, fmt="0.0%",
        note="大於 100%＝需求超過產能（營收由產能決定）")
    row("pgwe", "每 GW 台光電 AI 相關營收", "US$M/GW", {c_(y): f"={c_(y)}{R['ai']}/fx/{c_(y)}{R['gw']}*1000" for y in BU}, fmt="#,##0")
    sec_("C. 損益")
    row("gmn", "其他營收毛利率 ＝ 錨點毛利率 − 差距 × AI 占比（FY2026 校準）", "%", {"D": f"=a26_gp/D{R['rev']}-{P['d']}*D{R['aish']}"}, fmt="0.0%",
        note="Derived：讓 FY2026 毛利率等於錨點；AI 相關毛利率＝其他＋差距")
    row("gma", "AI 相關毛利率 ＝ 1 −（1 − 校準值）÷ 價格指數", "%", {"D": f"=D{R['gmn']}+{P['d']}", **{c_(y): f"=1-(1-$D${R['gmn'] + 1})/{c_(y)}{R['pp']}" for y in BU[1:]}}, fmt="0.0%",
        note="Derived：每張成本不變，價格升降全部落在毛利（組合升級的毛利效果未另計）")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"={c_(y)}{R['ai']}*{c_(y)}{R['gma']}+{c_(y)}{R['nai']}*$D${R['gmn']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in BU[1:]}}, fmt="#,##0.0")
    row("opx", "營業費用", "NT$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("oi", "營業利益", "NT$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.0")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("nop", "業外收益", "NT$B", {c_(y): "=nop_f" for y in BU[1:]}, fmt="#,##0.0", note="FY2025、FY2026 含在淨利錨點內")
    row("ni", "歸屬母公司淨利", "NT$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}+{c_(y)}{R['nop']})*nir_f" for y in BU[1:]}}, fmt="#,##0.0")
    row("nioi", "淨利 ÷ 營業利益", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['oi']}" for y in FYS}, fmt="0.0%")
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.0000")
    row("eps", "EPS", "NT$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 台光電損益（十億新台幣；EPS 新台幣；歸屬母公司）。FY2025A 實績、FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 相關營收（AI 伺服器＋交換器）", "NT$B", "#,##0.0"),
                                ("nai", "  非 AI 營收", "NT$B", "#,##0.0"), ("util", "  需求面 ÷ 產能上限", "%", "0.0%"), ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"),
                                ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("eps", "EPS", "NT$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            if key == "util" and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收 − 錨點；FY2026 模型 AI 相關 − 錨點；FY2026 毛利 − 錨點").font = BOLD; sr += 1
CHK = sr
PL.cell(sr, 1, "  情境 A：營收差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：AI 相關差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['ai']}-a26_ai"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['gma']}*D{PB['A']['ai']}+D{PB['A']['gmn']}*D{PB['A']['nai']}-a26_gp"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：產能上限 − 營收（NT$B；FY2026 校準恆等）"); PL[f"D{sr}"] = f"=D{PB['A']['rcap']}-D{PB['A']['rev']}"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：每 GW CCL 金額與份額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 每 GW 高速 CCL 金額（加速器／運算板、交換器板、UBB／中板）、台光電份額與產能（本專案權威，規則第 13 條）。部署 GW＝AI IT 資本支出（markets/network）÷ 每 GW IT 資本支出（Tokenomics，經橋接表）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("gw", "部署 GW（情境 A）", "GW", lambda y: f"=PL!{FC[y]}{A_['gw']}", "0.0", "AI IT 資本支出 ÷ 每 GW IT 資本支出"),
        ("vcmp", "每 GW CCL：加速器／運算板", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['vcmp']}", "#,##0", "2026 Assumed、之後 × 指數（Analogy）"),
        ("vsw", "每 GW CCL：交換器板", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['vsw']}", "#,##0", "Assumed"),
        ("vubb", "每 GW CCL：UBB／中板／背板", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['vubb']}", "#,##0", "Assumed；Rubin 中板、Rubin Ultra 正交背板"),
        ("pgw", "每 GW 高速 CCL 金額（合計，不含價格指數）", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgw']}", "#,##0", ""),
        ("prk", "每架等值 CCL 金額（以 VR200 每 GW 架數換算）", "US$K/架", lambda y: f"=PL!{FC[y]}{A_['pgw']}*1000/racks_vr", "#,##0.0", "對照 Morgan Stanley VR200 每架 US$18.75K（口徑較窄）"),
        ("mku", "AI 高速 CCL 市場（含價格指數）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}*PL!{FC[y]}{A_['pp']}", "#,##0.0", "含 ASIC 部署（假設每 GW 金額相同）"),
        ("sh", "台光電有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以公司 AI 相關占比校準"),
        ("aid", "台光電 AI 相關營收（需求面）", "US$B", lambda y: f"=PL!{FC[y]}{A_['aid']}/fx", "#,##0.00", ""),
        ("ai", "台光電 AI 相關營收（實際＝MIN 需求面、產能）", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx", "#,##0.00", ""),
        ("pgwe", "每 GW 台光電 AI 相關營收", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgwe']}", "#,##0", ""),
        ("cap", "平均月產能", "百萬張／月", lambda y: f"=PL!{FC[y]}{A_['cap']}", "0.00", "公司擴產計畫"),
        ("rps", "每張營收（含組合與價格指數）", "NT$／張", lambda y: f"=PL!$D${A_['rps0']}*PL!{FC[y]}{A_['mix']}*PL!{FC[y]}{A_['pp']}", "#,##0", "FY2026 校準"),
        ("util", "需求面 ÷ 產能上限", "%", lambda y: f"=PL!{FC[y]}{A_['util']}", "0.0%", "大於 100%＝產能受限"),
        ("oth", "其他廠商 AI 高速 CCL（台燿、斗山、生益、松下等）＝ 市場 − 台光電", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}*PL!{FC[y]}{A_['pp']}-PL!{FC[y]}{A_['ai']}/fx", "#,##0.0", "產能受限時，未滿足需求計入其他廠商")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：每 GW CCL 金額的絕對水準只影響 FY2026 校準出的份額（金額高 → 份額低），不影響營收；影響營收的是金額指數（世代升級）與份額相對值。2027–28 年需求面高於產能時，營收由月產能 × 每張營收決定，擴產進度與 M9 組合比需求更關鍵。").font = BOLD
UNOTE = ur
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」與十年 P/E 規則值（年底落後 P/E，基準為 FY2027 EPS）。本版未做先驗，無 Andy 合理 P/E 對照"
V["A2"] = "循環位置：台光電不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 NVIDIA、Arista、台達電；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-22 收盤）"); V["B6"] = "=px"; V["B6"].number_format = "#,##0"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "未來 12 月 EPS（模型 A）≈ ¼ FY2026＋¾ FY2027"); V["B11"] = f"=0.25*PL!D{SUMR[('A', 'eps')]}+0.75*PL!E{SUMR[('A', 'eps')]}"; V["B11"].number_format = "0.00"
V.cell(12, 1, "  現價 ÷ 未來 12 月 EPS（模型 A）"); V["B12"] = "=px/B11"; V["B12"].number_format = "0.0"
V.cell(13, 1, "近 12 月 P/E（Yahoo 股市、財報狗，10-22）"); V["B13"] = "=pe_ttm"; V["B13"].number_format = "0.0"
V["E8"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"
V.cell(15, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, 16, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 17):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(21, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 22, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 23 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{17 + j}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(26, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, 27, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數", "對照：十年 75 分位"])
vr = 28
for x in (0.5, 0.75, 1.0, 1.21, 1.5):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, "=pe_mid").number_format = "0.0"; V.cell(vr, 5, "=pe_trough").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數遠高於歷史中位數時，現價依賴「AI 期倍數高於 2017–2022 年」與 FY2028 以後的成長延續；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。"); WMT_NOTE = vr
assert vr < 50, vr
# 十年年底 P/E（2016–2025）：來源 1＝Yahoo Finance 12 月收盤 ÷ 當年 EPS（HiStock、nstock 年度 EPS）；2016 年 EPS 未找到
CLOSE = [90.0, 101.5, 65.7, 137.0, 157.0, 278.0, 171.0, 382.0, 618.0, 1645.0]
EPSY = [None, 8.73, 5.48, 10.14, 11.33, 16.5, 15.24, 16.35, 27.81, 41.67]
PE_D = [None if e is None else f"={c}/{e}" for c, e in zip(CLOSE, EPSY)]
PE_SA = [None] * 5 + [16.85, 11.22, 23.78, 22.27, 40.15]     # stockanalysis 年度 P/E（2021–2025）
PE_HS = [None] * 5 + [18.60, 10.94, 26.84, 23.93, 43.37]     # HiStock 12 月本益比（2021–2025）
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "年底收盤 ÷ 當年 EPS（Yahoo Finance 收盤、HiStock／nstock EPS）", "vals": PE_D, "tag": "Derived（2017–2025 共 9 年；2016 年 EPS：未找到）"},
    {"label": "stockanalysis 年度 P/E", "vals": PE_SA, "tag": "Interested-party（資料商；只有 2021–2025，只作核對，不採用）", "use": False},
    {"label": "HiStock 12 月本益比", "vals": PE_HS, "tag": "Interested-party（資料商；近四季 EPS 口徑；只有 2021–2025，不採用）", "use": False}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位。採用一個來源（年底收盤 ÷ 當年 EPS，Derived；2016 年 EPS 未找到，共 9 年）；stockanalysis、HiStock 只有 2021–2025，兩者與採用值同量級，只作核對。2017–2022 年台光電以手機、筆電 HDI 材料為主（約 11–17x），AI 前後倍數口徑不同，規則值可能低估 AI 期的合理倍數。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# ---------------- Val 第 0 節：目標價與上檔空間（v0.1.1 新增；放在既有空白列，不位移既有列） ----------------
TP0 = 35; TPR = {}
assert all(V.cell(rw, c).value is None for rw in range(TP0 - 1, TP0 + 9) for c in range(1, 9)), TP0
assert V["A3"].value is None
V["A3"] = f"目標價與上檔空間：見第 0 節（本頁第 {TP0} 列；v0.1.1 新增，為了不位移既有列而放在空白列）"
V.cell(TP0, 1, "舊法：規則值情境加權（參考；v0.1.1 的「0. 目標價與上檔空間」改名，不再是目標價；權重為 In 頁 tp_wA／tp_wB／tp_wC）").font = BOLD
for k_, (key_, lab_) in enumerate((("tp", "舊法隱含價（情境加權，參考）"), ("up", "舊法隱含價相對現價（÷ 現價 − 1）"),
        ("A", "情境 A 隱含價（樂觀）"), ("B", "情境 B 隱含價（基準）"), ("C", "情境 C 隱含價（下檔）"),
        ("w", "情境權重 A／B／C"), ("m", "舊法說明")), TP0 + 1):
    V.cell(k_, 1, lab_); TPR[key_] = k_
V.cell(TPR["tp"], 2, f"=tp_wA*B{TPR['A']}+tp_wB*B{TPR['B']}+tp_wC*B{TPR['C']}").number_format = "#,##0"
V.cell(TPR["tp"], 5, "＝權重 A × 情境 A 隱含價＋權重 B × 情境 B 隱含價＋權重 C × 情境 C 隱含價（NT$）")
V.cell(TPR["up"], 2, f"=B{TPR['tp']}/px-1").number_format = "0.0%"
V.cell(TPR["up"], 5, "相對「" + V["A6"].value + "」")
for s_ in SC:
    V.cell(TPR[s_], 2, f"=E{VAL[s_]}").number_format = "#,##0"
    V.cell(TPR[s_], 5, f"＝規則值隱含股價 E{VAL[s_]}（{SCN_T[s_]}）")
for j_, s_ in enumerate(SC):
    V.cell(TPR["w"], 2 + j_, f"=tp_w{s_}").number_format = "0%"
V.cell(TPR["w"], 5, "In 頁輸入（Assumed，Andy 可改）：B 欄＝情境 A、C 欄＝情境 B、D 欄＝情境 C")
V.cell(TPR["m"], 2, "P/E 規則值：各情境 FY2027 EPS × 2017–2025 年底 P/E 依循環位置取分位（見頂＝25 分位、中段＝中位數、谷底＝75 分位；年底收盤 ÷ 當年 EPS，一個來源、9 年）。隱含價約為 FY2027 年底價、未折現，作 12 個月目標價的近似；情境 A／B／C 即本模型三個顆數情境（A 最高、C 最低）。偏誤：十年倍數含 2017–2022 年手機、筆電 HDI 材料期低倍數（約 11–17x），目標價可能偏低。")
for c_ in range(1, 6):
    V.cell(TPR["tp"], c_).fill = KEY; V.cell(TPR["up"], c_).fill = KEY
V.cell(TPR["tp"], 1).font = BOLD
defname("old_rule_px", f"Val!$B${TPR['tp']}"); defname("old_rule_up", f"Val!$B${TPR['up']}")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額、每張營收、其他營收毛利率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 相關營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 相關", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 相關", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
def alias(new, base, suf, years=(27, 28, 29), sep="_"):
    for y in years:
        wb.defined_names[f"{new}_{y}"] = DefinedName(f"{new}_{y}", attr_text=wb.defined_names[f"{base}{sep}{y}{suf}"].attr_text)
for nm, base, suf in (("CXLO", "cmpx", "_lo"), ("CXHI", "cmpx", "_hi"), ("SXLO", "swx", "_lo"), ("SXHI", "swx", "_hi"), ("UXLO", "ubbx", "_lo"), ("UXHI", "ubbx", "_hi"),
                      ("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("PPLO", "pp", "_lo"), ("PPHI", "pp", "_hi"), ("CAPLO", "cap", "_lo"), ("CAPHI", "cap", "_hi"),
                      ("MXLO", "mix", "_lo"), ("MXHI", "mix", "_hi")):
    alias(nm, base, suf)
for nm, suf in (("GNLO", "_lo"), ("GNHI", "_hi")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}{y}"] = DefinedName(f"{nm}{y}", attr_text=wb.defined_names[f"g_nai{y}{suf}"].attr_text)
cases = [("每 GW 金額成長 低", "三項指數皆取低值（運算板 2027 1.20、交換器 1.15、中板 1.30）", dict(cmpx="CXLO", swx="SXLO", ubbx="UXLO")),
         ("每 GW 金額成長 高", "三項指數皆取高值（運算板 2027 1.80、交換器 1.70、中板 2.50）", dict(cmpx="CXHI", swx="SXHI", ubbx="UXHI")),
         ("中板／背板成長 低", "只動 UBB／中板指數（2027 1.30、2028 1.50）", dict(ubbx="UXLO")),
         ("中板／背板成長 高", "只動 UBB／中板指數（2027 2.50、2028 3.00）", dict(ubbx="UXHI")),
         ("份額流失（低）", "份額相對值 2027 0.85、2028 0.75、2029 0.70（台燿、斗山、生益追趕）", dict(rel="RLO")),
         ("份額上升（高）", "份額相對值 2027 1.05、2028–29 1.00", dict(rel="RHI")),
         ("價格回落快", "價格指數 2027 0.95、2028 0.88、2029 0.85", dict(pp="PPLO")),
         ("價格延續", "價格指數 2027 1.08、2028 1.05、2029 1.00", dict(pp="PPHI")),
         ("擴產延遲", "平均月產能 2027 7.0、2028 9.5、2029 11.1 百萬張", dict(cap="CAPLO")),
         ("擴產提前", "平均月產能 2027 8.2、2028 10.8、2029 13.0 百萬張", dict(cap="CAPHI")),
         ("每張組合升級 慢", "組合指數 2027 1.10、2028 1.15、2029 1.20", dict(mix="MXLO")),
         ("每張組合升級 快", "組合指數 2027 1.45、2028 1.70、2029 1.85", dict(mix="MXHI")),
         ("AI 部署 ×0.8", "FY2027 起 AI IT 資本支出 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("AI 毛利率差距 4pt", "AI 相關毛利率只高 4 個百分點", dict(d="d_ai_lo")), ("AI 毛利率差距 12pt", "高 12 個百分點", dict(d="d_ai_hi")),
         ("非 AI 成長 低", "2027–29 0%", dict(gn="GNLO")), ("非 AI 成長 高", "2027 +15%、2028 +12%、2029 +10%", dict(gn="GNHI"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "ai", 2027, "#,##0.0"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "ai", 2028, "#,##0.0"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
assert nn + 3 < 40
SENS_LAST = nn + 1; SENS_NOTE = nn + 3
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數情境 B、C 的差異從 FY2028 起。產能受限的年度，需求面輸入（每 GW 金額、份額、AI 部署）對營收的影響被產能上限截斷，只有產能與每張組合指數會改變營收；價格指數同時改變產能上限營收與 AI 毛利率（成本不變），所以對 EPS 的影響大於對營收。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 券商估計、研調份額與每架內容。只作事後對照，不為對上數字而調參數。券商多為單一來源；FY2028 共識（付費）未找到"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 券商估計（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 凱基（07-29）", "=b_rev26k", pl("rev", 2026), "凱基 2026-07-29", "3Q26 實際 608.9 億，高於凱基 556.0 億", fmt="#,##0.0")
rec("FY2026 EPS vs 凱基（07-29）", "=b_eps26k", pl("eps", 2026), "凱基 2026-07-29", "")
rec("FY2026 EPS vs 市場共識（07-29）", "=b_eps26c", pl("eps", 2026), "凱基報告所列共識", "")
rec("FY2026 EPS vs 本土投顧（10-01）", "=b_eps26l", pl("eps", 2026), "三立 2026-10-01", "3Q 營收公布前")
rec("FY2027 營收 vs 凱基", "=b_rev27k", pl("rev", 2027), "凱基 2026-07-29", "", fmt="#,##0.0")
rec("FY2027 營收 vs 本土投顧", "=b_rev27l", pl("rev", 2027), "三立 2026-10-01", "", fmt="#,##0.0")
rec("FY2027 EPS vs 凱基（07-29）", "=b_eps27k", pl("eps", 2027), "凱基 2026-07-29", "")
rec("FY2027 EPS vs 市場共識（07-29）", "=b_eps27c", pl("eps", 2027), "凱基報告所列共識", "")
rec("FY2027 EPS vs 本土投顧（10-01）", "=b_eps27l", pl("eps", 2027), "三立 2026-10-01", "三個來源區間 190.6–204.9")
rec("FY2028 營收 vs Morgan Stanley 隱含（複合成長 76%）", "=ms_rev28", pl("rev", 2028), "Morgan Stanley（Deep News 2026-09-23）", "由成長率推算，非原文數字", fmt="#,##0.0")
rec("FY2028 EPS vs Morgan Stanley 隱含（目標價 ÷ 20 倍）", "=ms_eps28", pl("eps", 2028), "同上", "由目標價推算，非原文數字")
RC.cell(rr_, 1, "B. 份額與每 GW 金額（機制檢查）").font = BOLD; rr_ += 1
rec("FY2026 台光電有效份額 vs Prismark 高速材料 2025", "=hs_sh25", lambda s: f"=PL!D{PB[s]['sh0']}", "Prismark（公司轉述）", "口徑不同：本模型分母只含 AI 部署（含 ASIC），不含非 AI 高速材料", fmt="0.0%", pt=True)
rec("FY2026 台光電有效份額 vs 公司自估（超過 50%）", "=hs_sh25c", lambda s: f"=PL!D{PB[s]['sh0']}", "公司（富果 2026-03-12）", "", fmt="0.0%", pt=True)
rec("CY2027 每 GW 金額指數（÷ CY2026）vs Morgan Stanley 每架金額換算", "=(ms_rk_vr*racks_vr)/(ms_rk_gb*racks_gb*(1-vr26)+ms_rk_vr*racks_vr*vr26)",
    lambda s: f"=PL!E{PB[s]['pgw']}/PL!D{PB[s]['pgw']}", "Morgan Stanley 每架 GB300 6,695、VR200 18,750 美元 × 每 GW 架數", "對照值以 2026 年 Rubin 占比混合", fmt="0.00")
rec("CY2027 每 GW 金額指數（÷ CY2026）vs 每架張數換算", "=(sht_vr*racks_vr)/(sht_gb*racks_gb*(1-vr26)+sht_vr*racks_vr*vr26)",
    lambda s: f"=PL!E{PB[s]['pgw']}/PL!D{PB[s]['pgw']}", "源達證券：每架 164 張 → 578 張", "張數不含材料等級升級（M8 → M9）", fmt="0.00")
rec("CY2027 每架等值 CCL（US$K，VR200 架數）vs Morgan Stanley VR200 每架", "=ms_rk_vr/1000", lambda s: f"=PL!E{PB[s]['pgw']}*1000/racks_vr", "Morgan Stanley", "本模型含交換器（scale-out）、ASIC 與 UBB；Morgan Stanley 口徑較窄（未載明），差距不具意義，不列", fmt="#,##0.0", gap=False)
rec("FY2027 每 GW 台光電 AI 相關營收（US$M）", None, pl("pgwe", 2027), "—", "", fmt="#,##0")
RC.cell(rr_, 1, "C. 產能與校準").font = BOLD; rr_ += 1
rec("FY2027 需求面 ÷ 產能上限", None, pl("util", 2027), "—", "大於 100%＝產能受限", fmt="0.0%")
rec("FY2028 需求面 ÷ 產能上限", None, pl("util", 2028), "—", "", fmt="0.0%")
rec("每張營收 FY2026（NT$；校準）", None, lambda s: f"=PL!D{PB[s]['rps0']}", "—", "FY2025：942.6 億 ÷（585 萬張 × 12）≈ NT$1,343（Q 頁）；漲價與 M8 組合使每張營收倍增", fmt="#,##0")
rec("其他營收毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "AI 相關毛利率＝此值＋d_ai", fmt="0.0%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "AI 占比上升使毛利率上升", fmt="0.0%", pt=True)
rec("4Q26 EPS（模型）", None, lambda s: "=Q!E15", "—", "本土投顧 3Q26 估 36.85（10-01）", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代「AI IT 資本支出 ÷ 每 GW IT 資本支出」"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝台光電法說，經富果、uanalyze、udn 轉述；季報細項取自凱基整理；券商與研調皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：4Q25 法說（富果 2026-03-12；FY2025、4Q25、產品組合、擴產、M8／M9、漲價、份額）", FG0312),
        ("公司：2Q26 法說（富果 2026-08-27；3Q26 指引、基礎設施占比、產能 2028 年 1,110 萬張）", FG0827),
        ("二手：uanalyze 2026-07-29（2Q26、1H26 營收、毛利率、淨利、EPS；Prismark 高速材料 38.6%）", UA0729),
        ("二手：uanalyze FY2025（EPS 41.67、營收 942.6 億、股利 25 元）", "https://uanalyze.com.tw/articles/5965248827"),
        ("二手：udn 1Q26（淨利 53.4 億、EPS 14.9）", "https://udn.com/news/story/7253/9472499"),
        ("月營收：TechNews 2026-10-05（9 月 215.32 億、前 9 月 1,412.27 億）", TN1005),
        ("券商：凱基 2026-07-29（2Q26 損益細項、3Q26 與 2026–27 預估、共識、產品組合、M7 以上占比）", KGI),
        ("券商：三立／鏡週刊 2026-10-01（本土投顧 2026–27 EPS、產能路徑、高階 CCL 供需）", SETN),
        ("券商：Morgan Stanley（經 Deep News／老虎 2026-09-23；CCL 每架金額、TAM、台光電成長率與目標價）", MS),
        ("研調：源達證券 2026-09-28（新浪轉載；每架 CCL 張數、板層數、Prismark 份額）", YD),
        ("資料商：Yahoo 股市、財報狗（10-22 收盤、近 12 月 P/E）", "https://tw.stock.yahoo.com/quote/2383.TW"),
        ("資料商：Yahoo Finance 月線收盤（2016–2025 年 12 月）", "https://finance.yahoo.com/quote/2383.TW/history/"),
        ("資料商：HiStock 年度 EPS、月本益比；nstock 季 EPS；stockanalysis 年度財報與 P/E", "https://histock.tw/stock/2383/%E6%9C%AC%E7%9B%8A%E6%AF%94"),
        ("資料商：玩股網（單季營業利益季增率，用於拆 1Q26 營業利益）", "https://www.wantgoo.com/stock/2383/enterprise-growth/operating-profit-margin"),
        ("市場層：markets/network、markets/chip_bridge、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：3Q26 財報與法說（10 月底）；10 月營收（11 月上旬）；法說原文（公司網站）；1H26 業外；2016 年 EPS；前瞻 P/E 十年歷史（Bloomberg）；每張 CCL 單價；FY2028 共識（付費）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("台光電（2383）收支與估值模型 v0.1.1（AI 半導體 Project：載板與 PCB 次產業；高速 CCL）　2026-10-25", True),
 ("版本：v0.1.1（2026-10-25：Val 新增第 0 節「目標價與上檔空間」、In 新增情境權重 tp_wA／tp_wB／tp_wC；其餘數值與 v0.1 相同）。v0.1（初版）：每 GW 高速 CCL 金額（加速器／運算板、交換器板、UBB／中板）、台光電份額與產能在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）；台燿、聯茂等 CCL 廠之後讀這條線。", False),
 ("模型命題：台光電的 AI 營收＝部署 GW（AI IT 資本支出 ÷ 每 GW IT 資本支出）× 每 GW 高速 CCL 金額（VR200 中板、M9 等級使每 GW 金額上升）× 台光電份額 × 價格指數，但受月產能 × 每張營收限制（全產全銷）。非 AI（通用伺服器、低軌衛星、手機、筆電、車用）另依成長率推。股價約 42 倍 FY2026、24 倍 FY2027 EPS（模型）。", False),
 ("核心問題：2027–28 年 AI 每 GW CCL 需求超過產能時，擴產（月產能 615 → 945 → 1,110 萬張）與 M9 組合升級能讓營收成長多少；價格與份額回落時 EPS 會少多少；現價隱含的路徑合不合理。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029（三年）。金額：十億新台幣；每 GW 金額以美元計、以匯率換算。口徑：歸屬母公司淨利。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("需求面：markets/network 的 AI IT 資本支出 × 情境顆數係數 ÷ 每 GW IT 資本支出（Tokenomics，經橋接表）＝部署 GW；× 每 GW CCL 金額（2026 年三項合計 US$350M，Assumed；之後 × 世代指數，Analogy／Assumed）× 匯率 × 有效份額（FY2026 以公司 AI 相關占比校準）× 份額相對值 × 價格指數（Assumed）。", False),
 ("產能面：平均月產能（公司擴產計畫，Derived）× 12 × 每張營收（FY2026 校準）× 組合指數（M8→M9，Assumed）× 價格指數。AI 相關營收＝MIN（需求面、產能上限 − 非 AI）。", False),
 ("毛利率：AI 相關比其他營收高 d_ai（Assumed），其他毛利率為 FY2026 殘差（Derived）；價格指數變動全部落在 AI 毛利。營業費用：彈性 phi（Derived）。淨利 ÷ 稅前 0.77（2Q26 實績）。", False),
 ("校準：FY2026 ＝ 1H26 實績＋3Q26 營收實績（月營收）＋假設（4Q 季增 6%、毛利率 35%、營業費用率 6.5%）。3Q26 法說截至 2026-10-25 尚未舉行。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：現價 ÷ FY2027、FY2028 EPS；十年 P/E 規則值（年底收盤 ÷ EPS，9 年）× FY2027 EPS；循環位置以本模型營收最高年判定（台光電不套設備股見頂年）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文＝法說（經轉述）與月營收；券商多為單一來源。", False),
 ("", False),
 ("已知限制：(1) 每 GW CCL 金額的 2026 年水準與世代指數皆為 Assumed／Analogy（研調每架金額與張數口徑不一）；水準只影響校準份額，指數影響成長；(2) 2027–28 年產能受限，營收主要由擴產進度與每張組合指數（Assumed）決定；(3) AI 相關占比（基礎設施 × AI 伺服器＋交換器）為公司區間的中點，交換器含非 AI；(4) 1H26 與 FY2025 業外、稅前未找到，以淨利 ÷ 營業利益涵蓋；(5) 十年 P/E 只有一個採用來源，2016 年缺，且 2017–2022 年為 AI 前低倍數；(6) 3Q26 財報與法說（10 月底）後更新錨點；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

V2["pe10"] = (list(zip(range(2016, 2026), PE_D)), "年底收盤 ÷ 當年 EPS；同規則值；2016 年 EPS 未找到")
V2["old_row"] = TP0
# ================= v0.2（2026-10-25）：PL 延伸到 FY2032、Val 第 0 節目標價（DCF 與 Comparable 加權）、Sens 目標價參數 =================
# 全部放在各頁既有內容之後，不位移既有儲存格（FY2025–FY2029 數值與 v0.1.1 相同）
FADEFILL = PatternFill("solid", fgColor="E7E6E6")
XY = (2025, 2026, 2027, 2028, 2029, 2030, 2031, 2032); XC = {y: "CDEFGHIJ"[i] for i, y in enumerate(XY)}; FADE = (2030, 2031, 2032)
K2 = V2["keys"]
for rw in range(4, IN.max_row + 1):
    if IN.cell(rw, 1).value in ("tp_wA", "tp_wB", "tp_wC", "tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "nim10_f"):
        IN.cell(rw, 5).fill = KEY
# --- PL 延伸表（三情境到 FY2032） ---
xr = PL.max_row + 3; XT = xr
PL.cell(xr, 1, "延伸：三情境到 FY2032（v0.2 新增）。FY2025–FY2029 引用上方推導區塊；FY2030 起為收斂法（灰底），非驅動因子推導：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 線性收斂到收斂目標，FY2032 達到").font = BOLD
xr += 1; XR = {}; XH = None
for s in SC:
    E = PB[s]; R = {}
    PL.cell(xr, 1, f"延伸｜{SCN_T[s]}").font = BOLD; xr += 1
    hdr(PL, xr, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    XH = XH or xr; xr += 1
    for key, lab, unit in (("rev", "營收（延伸）", V2["unit"]), ("g", "  營收年增", "%"), ("ni", "淨利（延伸）", V2["unit"]), ("nim", "  淨利率", "%"),
                           ("sh", "股數（延伸）", V2["shunit"]), ("eps", "EPS（延伸）", V2["cur"]), ("epsg", "  EPS 年增", "%"),
                           ("tg", "收斂目標：FY2032 營收年增", "%"), ("tm", "收斂目標：FY2032 淨利率", "%")):
        PL.cell(xr, 1, lab); PL.cell(xr, 2, unit); R[key] = xr; xr += 1
    G = lambda k: f"$G${R[k]}"
    for y in XY:
        c = XC[y]; pc = XC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={c}{E[K2['rev']]}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            PL[f"{c}{R['ni']}"] = f"={c}{E[K2['ni']]}"
            if y > 2025 or K2.get("sh25", True): PL[f"{c}{R['sh']}"] = f"={c}{E[K2['sh']]}"
            PL[f"{c}{R['eps']}"] = f"={c}{E[K2['eps']]}"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"={G('g')}+($C${R['tg']}-{G('g')})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"={G('nim')}+($C${R['tm']}-{G('nim')})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['sh']}"] = f"={pc}{R['sh']}*(1+{V2['shg']})"
            PL[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh']}"
        if y <= 2029: PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
        if y > 2025: PL[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("rev", "#,##0.0"), ("g", "0.0%"), ("ni", "#,##0.0"), ("nim", "0.0%"), ("sh", "0.0000"), ("eps", V2["epsfmt"]), ("epsg", "0.0%")):
            PL[f"{c}{R[key]}"].number_format = fm
            if y in FADE: PL[f"{c}{R[key]}"].fill = FADEFILL
    PL[f"C{R['tg']}"] = f"=fade_g{s}"; PL[f"C{R['tg']}"].number_format = "0.0%"
    PL[f"C{R['tm']}"] = {"A": f"=G{R['nim']}", "B": f"=G{R['nim']}*(1+nim10_f)/2", "C": f"=G{R['nim']}*nim10_f"}[s]; PL[f"C{R['tm']}"].number_format = "0.0%"
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導（上游市場層只到 2029）")
    PL.cell(R["tg"], 11, {"A": "In 頁 fade_gA＝tp_g＋2pt", "B": "In 頁 fade_gB＝tp_g", "C": "In 頁 fade_gC＝tp_g−2pt"}[s])
    PL.cell(R["tm"], 11, {"A": "＝本情境 FY2029 淨利率", "B": "＝FY2029 淨利率與十年平均替代值（FY2029 × nim10_f）的中點", "C": "＝十年平均替代值（FY2029 × nim10_f；十年平均未找到）"}[s])
    PL.cell(R["sh"], 11, f"FY2030 起股數年變動＝{V2['shg']}")
    XR[s] = R; xr += 1
PL.column_dimensions["I"].width = 14; PL.column_dimensions["J"].width = 14; PL.column_dimensions["K"].width = 60
# --- Val 第 0 節：目標價 ---
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
t0 = V.max_row + 3; T0 = t0
V.cell(t0, 1, "0. 目標價（DCF 與 Comparable 加權；v0.2 新增，放在本頁既有內容之後以免位移既有儲存格；三情境欄 B／C／D＝A 順利／B 平平／C 不如預期）").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["h1"] = t0 + 2
r_ = t0 + 3
V.cell(r_, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{r_}"] = "PL 頁延伸表；FY2030 起為收斂法"; T["eps_h"] = r_; r_ += 1
EPSR = {}
for y in YRS6:
    tset(r_, f"eps{y}", f"  FY{y}", [f"=PL!{XC[y]}{XR[s]['eps']}" for s in SC], V2["epsfmt"]); EPSR[y] = r_; r_ += 1
V.cell(r_, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = r_; r_ += 1
tset(r_, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）"); RV = r_; r_ += 1
CFR = {}
for y in YRS6:
    tset(r_, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], V2["epsfmt"]); CFR[y] = r_; r_ += 1
def dcf_formula(c, r="tp_r", g="tp_g", reinv=None):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    reinv = reinv or f"Val!$B${RV}"
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
c1, c6 = CFR[2027], CFR[2032]
tset(r_, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], V2["pxfmt"],
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027"); r_ += 1
tset(r_, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{T['dcfv']}*(1+tp_r)" for c in "BCD"], V2["pxfmt"]); r_ += 1
tset(r_, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{T['dcf12']}+tp_wB*C{T['dcf12']}+tp_wC*D{T['dcf12']}", None, None], V2["pxfmt"], "Σ tp_wA／wB／wC × 情境價值 ×(1＋tp_r)；權重見 In 頁"); r_ += 2
V.cell(r_, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; r_ += 1
hdr(V, r_, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["ph"] = r_; r_ += 1
P1 = r_
for nm, tk, pe, dt in V2["peers"]:
    V.cell(r_, 1, nm); V.cell(r_, 2, tk); cc = V.cell(r_, 3, pe); cc.font = BLUE; cc.number_format = "0.0"
    V.cell(r_, 4, dt); V.cell(r_, 5, "stockanalysis.com 個股頁 Forward PE（分析師共識；價格日期如左）"); V.cell(r_, 6, "Interested-party（賣方共識）"); r_ += 1
P2 = r_ - 1; T["p1"], T["p2"] = P1, P2
tset(r_, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{P1}:C{P2})", None, None], "0.0", V2["peer_note"]); r_ += 1
w1, w2 = V2["fwd_w"]
tset(r_, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*({w1}*{c}{EPSR[2027]}+{w2}*{c}{EPSR[2028]})" for s, c in zip(SC, "BCD")), None, None], V2["epsfmt"], V2["fwd_note"]); r_ += 1
tset(r_, "comptp", "Comparable 目標價", [f"=B{T['pmed']}*B{T['feps']}", None, None], V2["pxfmt"], "同業中位數 × 上列 EPS"); r_ += 2
V.cell(r_, 1, "0.3 結論").font = BOLD; r_ += 1
hdr(V, r_, ["項目", "數值", "", "", "說明"]); T["h3"] = r_; r_ += 1
tset(r_, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], V2["pxfmt"], "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁"); r_ += 1
tset(r_, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], V2["pxfmt"], "引用 0.2"); r_ += 1
tset(r_, "cons", "市場共識目標價", ["=cons_tp", None, None], V2["pxfmt"], V2["cons_tp"][1]); r_ += 1
tset(r_, "upw", "上檔空間：加權目標價", [f"=B{T['tpw']}/px-1", None, None], "0.0%", V2["px_note"]); r_ += 1
tset(r_, "upc", "上檔空間：Comparable 目標價", [f"=B{T['comp2']}/px-1", None, None], "0.0%"); r_ += 1
tset(r_, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS 見 In 頁 cons_feps 與 Recon"); r_ += 1
tset(r_, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%"); r_ += 1
V[f"E{T['gmul']}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(r_, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）"); r_ += 1
PEROW = r_ + 3
PE10 = f"$B${PEROW + 1}:$K${PEROW + 1}"
tset(r_, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
     [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E 分布（下方列）的分位；低於最小值記 0%、高於最大值記 100%；只作檢查，不進目標價"); r_ += 1
V.cell(r_, 1, "目標價方法"); V[f"B{r_}"] = V2["method"]; T["m"] = r_; r_ += 1
V.cell(PEROW, 1, "十年年底 P/E（歷史分位檢查用；" + V2["pe10"][1] + "）").font = BOLD
for j, (yy, vv) in enumerate(V2["pe10"][0]):
    V.cell(PEROW, 2 + j, str(yy))
    cc = V.cell(PEROW + 1, 2 + j, vv); cc.font = BLUE; cc.number_format = "0.0"
V.cell(PEROW + 1, 1, "年底 P/E"); T["pe10h"] = PEROW; T["pe10"] = PEROW + 1
for k_ in ("dcftp", "comptp", "tpw"):
    V[f"B{T[k_]}"].fill = KEY
V.cell(T["tpw"], 1).font = BOLD
defname("tp_final", f"Val!$B${T['tpw']}"); defname("tp_dcf", f"Val!$B${T['dcftp']}"); defname("tp_comp", f"Val!$B${T['comptp']}")
V["A3"] = f"目標價：見第 0 節（本頁第 {T0} 列；v0.2 新增，DCF 與 Comparable 加權）。第 {V2['old_row']} 列「舊法：規則值情境加權（參考）」不再是目標價，只作參考"
for c in "IJK": V.column_dimensions[c].width = 10
# --- Sens：目標價參數 ---
TPS0 = SE.max_row + 3
SE.cell(TPS0, 1, "目標價參數（v0.2）：一次動一個，對加權目標價的影響（只改 DCF 參數；PL 收斂目標 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = V2["pxfmt"]
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "tp_r 10%、tp_g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "tp_r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "tp_r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "tp_g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "tp_g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.4", "tp_w_dcf 0.4", dict(w="0.4")), ("DCF 權重 0.8", "tp_w_dcf 0.8", dict(w="0.8"))]
for k_, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k_, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
# --- Recon：目標價對帳 ---
rq = RC.max_row + 2
RC.cell(rq, 1, "E. 目標價對帳（v0.2；模型欄為情境加權值，非情境 A）").font = BOLD; rq += 1
RECV2 = rq
for lab, ref, mdl, src in (("12 個月後的前瞻 EPS（情境加權）vs 共識前瞻 EPS", "=cons_feps", f"=Val!B{T['feps']}", V2["cons_feps"][1]),
                           ("目標價（DCF 與 Comparable 加權）vs 市場共識目標價", "=cons_tp", f"=Val!B{T['tpw']}", V2["cons_tp"][1]),
                           ("Comparable 目標價 vs 市場共識目標價", "=cons_tp", f"=Val!B{T['comptp']}", "同上")):
    RC.cell(rq, 1, lab); RC.cell(rq, 2, ref).number_format = V2["epsfmt"]; RC.cell(rq, 3, mdl).number_format = V2["epsfmt"]
    RC.cell(rq, 4, f"=C{rq}/B{rq}-1").number_format = "0.0%"; RC.cell(rq, 9, src); RC.cell(rq, 10, "共識只作對帳，不作校準"); rq += 1
# --- README ---
for rw in range(1, RD.max_row + 1):
    v_ = RD.cell(rw, 1).value
    if isinstance(v_, str) and rw == 1: RD.cell(rw, 1, v_.replace("v0.1.1", "v0.2"))
    elif isinstance(v_, str) and v_.startswith("估值："): RD.cell(rw, 1, V2["readme_val"] + "（v0.1.1 原估值說明，保留作參考：" + v_[3:] + "）")
    elif isinstance(v_, str) and v_.startswith("版本："): RD.cell(rw, 1, "版本：v0.2（2026-10-25：目標價改為 DCF 與 Comparable 加權、三情境延伸到 FY2032；FY2025–FY2029 數值不變）。前版：" + v_[3:])
rn = RD.max_row + 1
for t_ in V2["readme_add"]:
    RD.cell(rn, 1, t_); rn += 1
if os.environ.get("V2_ROWMAP"):
    import json as _json
    _json.dump({"XT": XT, "XH": XH, "XR": XR, "T": T, "EPSR": EPSR, "CFR": CFR, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "RECV2": RECV2},
               open(os.environ["V2_ROWMAP"], "w"), ensure_ascii=False, indent=1)

# ---------------- v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00／14:17）----------------
# 新列一律放在 Val、README 工作表既有內容之後，不位移既有列；既有數值不變
VER_PREV, VER_NOW = 'v0.2', 'v0.2.1'
OBS = [
 "目標價 NT$5,477（DCF 4,252、Comparable 6,703 各半），比現價 5,070（2026-10-22）高 8.0%，級距 0%–25%；Comparable 單獨為 +32.2%。DCF 比現價低 16.1%、Comparable 高 32.2%，兩者相抵；本模型前瞻 EPS 254.18 比共識前瞻 EPS 168.32 高 51.0%（共識取自 2026-09-04 快取、期間早約一年），同業（台燿、聯茂、金像電）前瞻 P/E 中位數 26.4 倍，共識目標價 7,803 隱含 46.4 倍。",
 "價格是最敏感的變數：價格回落快（2028 年指數 0.88）使 FY2028 EPS −26.2%、價格延續 +23.9%；每張組合升級慢 −19.5%，顆數情境 C −30.7%。",
 "FY2027 EPS 三情境同為 209.59，FY2028 起分歧（299.15／263.73／207.44），DCF 價格帶 2,691～5,428。下一個檢驗點是 3Q26 財報與法說（預計 2026-10 底，確切日期未找到）。"
]   # 重點觀察（數字取自重算後的 Val／PL／Sens；與 README.md 相同文字）
UPW_ROW = next(r_ for r_ in range(1, V.max_row + 1) if V.cell(r_, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 1; REC_ROW = GRADE_ROW + 1
_u = f"B{UPW_ROW}"
V.cell(GRADE_ROW, 1, "上檔空間級距")
V[f"B{GRADE_ROW}"] = f'=IF({_u}>0.5,"大於 50%",IF({_u}>0.25,"25%–50%",IF({_u}>=0,"0%–25%",IF({_u}>=-0.25,"0%～−25%",IF({_u}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"B{GRADE_ROW}"].number_format = "@"
V[f"E{GRADE_ROW}"] = f"以「上檔空間：加權目標價」（第 {UPW_ROW} 列）分六級（Andy 2026-10-25 14:00）"
V.cell(REC_ROW, 1, "投資建議")
V[f"B{REC_ROW}"] = f'=IF({_u}>0.5,"強力買進",IF({_u}>0.25,"買進",IF({_u}>=-0.25,"中立",IF({_u}>=-0.5,"放空","強力放空"))))'
V[f"B{REC_ROW}"].number_format = "@"
V[f"E{REC_ROW}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）"
_t = RD["A1"].value
assert VER_PREV in _t and VER_NOW not in _t, _t
RD["A1"] = _t.replace(VER_PREV, VER_NOW, 1)
for r_ in range(2, RD.max_row + 1):
    _v = RD.cell(r_, 1).value
    if isinstance(_v, str) and _v.startswith("版本："):
        RD.cell(r_, 1, "版本：" + VER_NOW + "（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）。前版：" + _v[3:])
        break
rn = RD.max_row + 1
RD.cell(rn, 1, VER_NOW + "：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁最後新增「上檔空間級距」「投資建議」兩列（第 " + str(GRADE_ROW) + "、" + str(REC_ROW) + " 列，依「上檔空間：加權目標價」的固定規則）；重點觀察見下方三列。")
RD.cell(rn + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"')
for k_, o_ in enumerate(OBS, 1):
    RD.cell(rn + 1 + k_, 1, f"重點觀察 {k_}：" + o_)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("EMC_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "SB": list(SB), "SENS_LAST": SENS_LAST,
               "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "WMT_NOTE": WMT_NOTE, "UNOTE": UNOTE, "block_end": block_end},
              open(os.environ["EMC_ROWMAP"], "w"), ensure_ascii=False, indent=1)
