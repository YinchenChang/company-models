# Amazon（AMZN）個股模型 v0.3.1（命題＋營收分部架構；工作單 mag/docs/workorders/20261027_MAG_v0.3_個股模型重建_分部架構.md；
#   v0.3.1 季度滾動：mag/docs/workorders/20261029_MAG_v0.3.1_季度滾動_2026Q3.md 第 1–4 節）
# 固定規格：docs/plan/個股模型原則_命題加營收分部.md（命題模組＋官方營收分部 → 損益表；第 7 節讀法 2、第 8 節 In 頁、第 9 節敏感度覆蓋）。
# 用法：python3 amazon/build_amazon.py [輸出.xlsx] [--set 鍵=值 ...] [--sens sens.json] [--q3v03 q3.json]
#       （預設輸出 amazon/dist/20261031_Amazon個股模型_v0.3.1.xlsx；amazon/make_v031.sh 一次跑完 建置 → 重算 → 核對 → HTML）
#       再以 LibreOffice headless 重算：soffice --headless --calc --convert-to xlsx:"Calc MS Excel 2007 XML" --outdir <dir> <raw.xlsx>
# 全部活公式；藍字＝輸入、黑字＝公式。金額 US$bn（每股 US$）；MW＝IT 關鍵電力 MW（MW-IT）。
# 欄位（模型頁共用）：C FY2024A、D FY2025A、E 9M26A、F Q4 26E、G FY2026E（＝E＋F）、H–M FY2027–FY2032、N 9M25A、O Q4 25A。資料截至 2026-09-30。
# 工作表（順序固定）：README → In → Q → Thesis → Seg → PL → CF → Val → Recon → Sens → Sources
import argparse
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
AP = argparse.ArgumentParser()
AP.add_argument("out", nargs="?", default=os.path.join(HERE, "dist", "20261031_Amazon個股模型_v0.3.1.xlsx"))
AP.add_argument("--set", action="append", default=[], help="覆寫 In 頁基準值：鍵=值（敏感度與拆解重算用；值可為數字或以 = 開頭的公式）")
AP.add_argument("--sens", help="敏感度與拆解鏈的重算結果（amazon/sens_v031.py 產生）")
AP.add_argument("--q3v03", help="變動拆解用：以 v0.3 對 2026Q3 的模型值取代 2026Q3 實績（JSON：{事實 id: 值}；amazon/q3cf_v03.py 產生）")
ARGS = AP.parse_args()
OUT = ARGS.out
OVR = {}
for _s in ARGS.set:
    _k, _v = _s.split("=", 1)
    try:
        OVR[_k] = float(_v)
    except ValueError:
        OVR[_k] = _v if _v.startswith("=") else "=" + _v
SENS = json.load(open(ARGS.sens, encoding="utf-8")) if ARGS.sens else {}

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); CHK = PatternFill("solid", fgColor="E2EFDA"); SUBF = PatternFill("solid", fgColor="F2F2F2")
SC = ("A", "B", "C")
SCN = {"A": "A（樂觀）", "B": "B（基準）", "C": "C（保守）"}

# ---------------- 事實總帳（data/amazon_facts_20261008.json：至 2026Q2；data/amazon_facts_2026q3.json：2026Q3 新聞稿，v0.3.1） ----------------
FACTS = {}
for g in json.load(open(os.path.join(HERE, "data", "amazon_facts_20261008.json"), encoding="utf-8"))["groups"].values():
    for f in g:
        FACTS[f["id"]] = f
for f in json.load(open(os.path.join(HERE, "data", "amazon_facts_2026q3.json"), encoding="utf-8"))["items"]:
    FACTS[f["id"]] = f
if ARGS.q3v03:
    for k, v in json.load(open(ARGS.q3v03, encoding="utf-8"))["facts"].items():
        FACTS.setdefault(k, {"id": k})
        FACTS[k] = dict(FACTS[k], value=v)
CONS = json.load(open(os.path.join(HERE, "data", "consensus_amzn_20261008.json"), encoding="utf-8"))
SNAP = json.load(open(os.path.join(HERE, "data", "tokenomics_snapshot_v21.7.json"), encoding="utf-8"))
GEN = {g["code"]: g["name"] for g in SNAP["source"]["generations"]}
_PJ = os.path.join(HERE, "data", "pace_v031.json")
PACE = json.load(open(_PJ, encoding="utf-8")) if os.path.exists(_PJ) else {}


def tk(name, code=None, case="基準"):
    v = SNAP["items"][name]["values"]
    if code is None:
        return v
    x = v[GEN[code]]
    return x[case] if isinstance(x, dict) else x


def fv(fid, scale=1000.0):
    """事實值（USD M → US$bn）。"""
    return FACTS[fid]["value"] / scale


wb = openpyxl.Workbook()


def hdr(ws, row, labels, c0=1):
    for j, h in enumerate(labels, c0):
        if h is None:
            continue
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR


def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ---------------- 來源編號（Sources 頁） ----------------
SRC = []          # (編號, 文件, 日期, 網址, 標記)
SRCID = {}


def src(doc, date, url, tag):
    k = (doc, url)
    if k not in SRCID:
        SRCID[k] = f"S{len(SRC) + 1:02d}"
        SRC.append((SRCID[k], doc, date, url, tag))
    return SRCID[k]


S_10K = src("Form 10-K 2025（期末 2025-12-31，經查核）", "2026-02-06", "https://www.sec.gov/Archives/edgar/data/1018724/000101872426000004/amzn-20251231.htm", "Verified")
S_10Q = src("Form 10-Q 2026Q2（期末 2026-06-30，核閱未查核）", "2026-07-31", "https://www.sec.gov/Archives/edgar/data/1018724/000101872426000026/amzn-20260630.htm", "Verified")
S_10Q3 = src("Form 10-Q 2025Q3（分部 D&A 前九個月）", "2025-10-31", "https://www.sec.gov/Archives/edgar/data/1018724/000101872425000123/amzn-20250930.htm", "Verified")
S_PR26Q2 = src("2026Q2 財報新聞稿（8-K 附件 99.1；逐季補充表、指引）", "2026-07-30", "https://www.sec.gov/Archives/edgar/data/1018724/000101872426000024/amzn-20260630xex991.htm", "Interested-party")
S_PR26Q1 = src("2026Q1 財報新聞稿（8-K 附件 99.1）", "2026-04-29", "https://www.sec.gov/Archives/edgar/data/1018724/000101872426000012/amzn-20260331xex991.htm", "Interested-party")
S_PR25Q4 = src("2025Q4 財報新聞稿（8-K 附件 99.1）", "2026-02-05", "https://www.sec.gov/Archives/edgar/data/1018724/000101872426000002/amzn-20251231xex991.htm", "Interested-party")
S_PR25Q3 = src("2025Q3 財報新聞稿（8-K 附件 99.1）", "2025-10-30", "https://www.sec.gov/Archives/edgar/data/1018724/000101872425000121/amzn-20250930xex991.htm", "Interested-party")
S_PR24Q4 = src("2024Q4 財報新聞稿（8-K 附件 99.1；逐季補充表）", "2025-02-06", "https://www.sec.gov/Archives/edgar/data/1018724/000101872425000002/amzn-20241231xex991.htm", "Interested-party")
S_PR24Q3 = src("2024Q3 財報新聞稿（8-K 附件 99.1）", "2024-10-31", "https://www.sec.gov/Archives/edgar/data/1018724/000101872424000158/amzn-20240930xex991.htm", "Interested-party")
S_XBRL = src("SEC XBRL companyfacts（稀釋加權股數）", "2026-10-08", "https://data.sec.gov/api/xbrl/companyfacts/CIK0001018724.json", "Verified")
S_CALL = src("2026Q2 法說會逐字稿（二手轉錄：Webull；2026 現金資本支出約 220）", "2026-07-30", "https://www.webull.com/news/15364638750483456", "Interested-party")
S_424 = src("424B5 英鎊票據公開說明書補充（£3.25B）", "2026-09-11", "https://www.sec.gov/Archives/edgar/data/1018724/000110465926107122/tm2624614-3_424b5.htm", "Interested-party")
S_PR26Q3 = src("2026Q3 財報新聞稿（PDF 逐列轉錄 mag/extract/docs/2026Q3/amazon_2026Q3_新聞稿.md；IR 網頁第二讀 22 項一致）", "2026-10-29",
               "https://s2.q4cdn.com/299287126/files/doc_earnings/2026/q3/earnings-result/AMZN-Q3-2026-Earnings-Release.pdf", "Interested-party")
S_CALL3 = src("2026Q3 法說會逐字稿（不入 repo；只引用數字與短語：2026 現金資本支出約 220、2027「meaningfully more」、backlog 578、run-rate 184、AWS 2027–2028 各 >40%）", "2026-10-29",
              "https://ir.aboutamazon.com/（法說 webcast；逐字稿由 Andy 提供）", "Interested-party")
S_PX = src("Motley Fool／Yahoo Finance「Stock Market Today, Oct. 30」：AMZN 收盤 240.35（−1.41%）", "2026-10-30", "https://finance.yahoo.com/markets/stocks/articles/stock-market-today-oct-30-213502210.html", "Verified")
S_PXY = src("MarketChameleon AMZN 概覽：2026-10-30 收盤 240.35", "2026-10-30", "https://marketchameleon.com/Overview/AMZN", "Verified")
S_FC = src("BigGo Finance（Koyfin 彙整：60 位平均目標價 332.41、區間 230–405；TD Cowen 340、UBS 352、Mizuho 335 等財報後調整）", "2026-10-30", "https://finance.biggo.com/news/41346e0c-47f7-4ab6-adb0-e12aeffbd88f", "Interested-party")
S_MB = src("MarketBeat 預估頁（60 位共識目標價 325.12；最近更新 2026-10-29 收盤後）", "2026-10-29", "https://www.marketbeat.com/stocks/NASDAQ/AMZN/Forecast/", "Interested-party")
S_SP = src("StockAnalysis 預估頁（S&P Global：營收與調整後 EPS 共識；財報前）", "2026-10-26", "https://stockanalysis.com/stocks/amzn/forecast/", "Interested-party")
S_MS = src("MarketScreener 年度共識（S&P Capital IQ 系；EBITDA、EBIT、淨利、GAAP EPS、資本支出、FCF；財報前，本輪未能更新）", "2026-10-08", "https://www.marketscreener.com/quote/stock/AMAZON-COM-INC-12864605/finances/", "Interested-party")
S_TK = src("Tokenomics v21.7 快照（amazon/data/tokenomics_snapshot_v21.7.json；model/CURRENT＝20261025_Tokenomics_v21.7.xlsx，master 90f7deb）", "2026-10-31", "amazon/data/tokenomics_snapshot_v21.7.json", "Derived")
S_V03 = src("Amazon v0.3 成品與報告（加權目標價 215.58；拆解起點）", "2026-10-27", "amazon/dist/_待刪除/20261027_Amazon個股模型_v0.3.xlsx；amazon/docs/reports/20261027_amazon_v0.3_分部架構.md", "Derived")
S_KEV = src("k 證據（amazon/data/kevidence_amazon_20261008.json；Silicon Data H100 指數、AWS p5 牌價）", "2026-10-08", "amazon/data/kevidence_amazon_20261008.json", "Verified")
S_EPOCH = src("Epoch AI 資料中心目錄（New Carlisle、Ridgeland、Madison）", "2026-10-05", "https://epoch.ai/data/ai-data-centers/directory/", "Analogy")
S_V01A = src("v0.1a 資料報告（AI 在役 MW、對外比例、世代占比的推導）", "2026-10-08", "amazon/docs/reports/20261008_amazon_v0.1a_資料.md", "Derived")
S_V022 = src("v0.2.2 交接檔（MW 速度、k 規格、WACC、持股、同業倍數；本版沿用）", "2026-10-25", "amazon/docs/handoff/20261025_Amazon收支模型_交接檔_v0_2_2.md；引擎 amazon/legacy_v0.2/", "Derived")
S_PEERC = src("共用非 AI 雲端同業檔（ORCL、IBM、SAP、CRM、NOW、AKAM、EQIX、DOCN；StockAnalysis EV ÷ MarketScreener NTM EBITDA）", "2026-10-07", "amazon/data/peers_cloud_20261008.json", "Derived")
S_PEERA = src("共用數位廣告同業檔（META、PINS、TTD、APP、RDDT、SNAP）", "2026-10-07", "amazon/data/peers_ads_20261009.json", "Derived")
S_PEERR = src("零售同業（WMT、COST、TGT；事實總帳 val.peer.retail.*）", "2026-10-07", "amazon/data/amazon_facts_20261008.json", "Derived")
S_ANT = src("Anthropic Series H 投後估值 965B（W.Media）", "2026-06-01", "https://w.media/anthropic-valuation-touches-us-965-billion-surpassing-openai/", "Interested-party")
S_FORT = src("Fortune：Amazon 持有 Anthropic 約 mid-to-high teens", "2026-06-04", "https://www.fortune.com/2026/06/04/amazon-google-billions-anthropic-ipo/", "Analogy")
S_OAI = src("GeekWire：Amazon 投資 OpenAI 50B、投前 730B", "2026-02-27", "https://www.geekwire.com/2026/amazon-invests-50b-in-openai-deepens-aws-partnership-with-expanded-100b-cloud-deal/", "Interested-party")
S_UST = src("美國財政部每日殖利率曲線（10 年 4.99%，2026-10-28；Alphabet v0.3.1 同日同源；Mortgage News Daily 10-28 4.996% 佐證）", "2026-10-28", "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?type=daily_treasury_par_yield_curve_rates", "Verified")
S_BLOOM = src("Bloomberg／SourceMaterial：第三方託管約占 Amazon 運算力五分之一（2024）", "2025-11-24", "amazon/data/amazon_facts_20261008.json → mw.colo.share", "Interested-party")
S_MSTAN = src("Morgan Stanley（Brian Nowak）：AWS 2026 新增 6 GW、2027 新增 8 GW（Invezz 轉述）", "2026-08-18", "https://invezz.com/news/2026/08/18/could-amazon-stock-double-by-end-of-next-year-heres-what-morgan-stanley-says/", "Interested-party")

# =====================================================================================
# In
# =====================================================================================
IN = wb.active; IN.title = "In"
IN["A1"] = ("In — 輸入（藍字；基準欄 E 進主模型；D／F＝敏感度低／高，只用於 Sens 頁〔一次動一個參數〕、不進目標價；情境 A／B／C 的驅動見本頁最後一節「J. 情境驅動對照」〔v0.3.1〕）。"
            "金額 US$bn、每股 US$、MW＝IT 關鍵電力 MW。每列一個鍵（定義名稱＝鍵）；Analogy、Assumed 一律給低／高區間")
IN["A2"] = ("標記：Verified／Interested-party／Analogy／Assumed／Derived。「沿用 v0.2.2」＝v0.2.2 已核准的參數與證據（k 規格 kspec.py、對外比例），出處見 legacy_v0.2/ 與 docs/handoff。"
            "不以公司 AWS 營收、backlog、run-rate 反推 k 或每 MW 價格（共同規則第 4 節）；公司說法只作 Recon 驗證列。SemiAnalysis 未使用。資料截至 2026-09-30（2026Q3 新聞稿；10-Q 未取得）")
hdr(IN, 3, ["鍵", "項目", "單位", "敏感度低", "基準", "敏感度高", "標記", "來源", "說明"])
ROWS = []
def sec(t): ROWS.append(("sec", t))
def inp(*a): ROWS.append(a)
KEYFILL = set()
NM = "只作對照，不進模型"

sec("A. 價格、股數與共識目標價")
inp("px", "AMZN 現價（2026-10-30 收盤）", "US$", None, 240.35, None, "Verified", f"{S_PX}；{S_PXY}", "10-29 收盤 243.79（財報日，−4.3%）；只影響上檔空間與評等")
inp("px_val", "WACC 權益權重用股價（最近收盤；評價日 2026-09-30 收盤本輪未取得）", "US$", None, 240.35, None, "Verified", f"{S_PX}", "與現價分開：現價變動只影響上檔空間（拆解 (e)），本列變動歸 (f)")
inp("sh_out", "流通股數（2026-09-30，資產負債表股本說明）", "十億股", None, round(FACTS["sh.out.20260930"]["value"] / 1000, 6), None, "Interested-party", S_PR26Q3, "10-Q 封面股數未取得；Supplemental 同值 10,811 百萬股")
inp("sh_awd", "未歸屬股份獎勵（2026-09-30）", "十億股", None, fv("sh.awards.2026q3"), None, "Interested-party", S_PR26Q3, "223 百萬股（Supplemental）")
inp("sh_val", "分部加總用稀釋股數＝流通股數＋未歸屬獎勵（倍數為 SBC 加回口徑，故計入稀釋）", "十億股", None, "=sh_out+sh_awd", None, "Derived", "修正一（F2）", "DCF 視 SBC 為現金成本，改用流通股數 sh_out，不另計稀釋")
inp("cons_tp", "市場共識目標價（Koyfin 60 位平均，2026-10-30 財報後）", "US$", 230, 332.41, 405, "Interested-party", f"{S_FC}；{S_MB}", "MarketBeat 60 位 325.12（10-29）；區間為最低／最高目標價；" + NM)

sec("B. 折現率（CAPM）")
inp("rf", "無風險利率（10 年美債，2026-10-28）", "%", None, 0.0499, None, "Verified", S_UST, "v0.3 為 5.00%（10-26）")
inp("beta_sa", "原始 Beta：StockAnalysis 5 年月資料（Yahoo 同值，推測同源）", "x", None, 1.46, None, "Verified", "沿用 v0.2.2（事實總帳 val.beta.sa）", "")
inp("beta_gf", "原始 Beta：GuruFocus 3 年", "x", None, 1.567, None, "Verified", "沿用 v0.2.2（事實總帳 val.beta.gurufocus）", "不同資料商、不同期間（C13 獨立來源）")
inp("beta", "Beta（Blume 調整 0.67 × 原始＋0.33，兩個獨立來源的中位數）", "x", "=MIN(0.67*beta_sa+0.33,0.67*beta_gf+0.33)", "=MEDIAN(0.67*beta_sa+0.33,0.67*beta_gf+0.33)", "=MAX(0.67*beta_sa+0.33,0.67*beta_gf+0.33)", "Derived", "修正一（F4）", "低／高＝單一來源")
inp("erp", "股權風險溢酬", "%", 0.04, 0.05, 0.06, "Assumed", "沿用 v0.2.2（對照表預設 10）", "")
inp("kd_new", "新增借款利率（2026-07 新發 2066 票據有效利率；只用於利息費用）", "%", None, 0.0633, None, "Interested-party", S_10Q, "期後事項：固定票面 4.60%–6.25%、有效 4.67%–6.33%")
inp("kd_spr", "AA 級公司債約 10 年期利差（相對 10 年美債）", "%", 0.005, 0.0065, 0.01, "Assumed", "修正一（F4）：kd＝約 10 年期公司債殖利率", "Amazon 各期票據的 10 年期殖利率：未找到")
inp("kd", "稅前債務成本（WACC 用）＝rf＋利差", "%", None, "=rf+kd_spr", None, "Derived", "修正一（F4）", "")
inp("tax", "稅率（2023–2025 有效稅率平均）", "%", 0.135, "=AVERAGE(«Q.etr.C»,«Q.etr.D»,«Q.etr.E»)", 0.23, "Derived", S_10K, "18.96%、13.50%、19.61%；9M26 有效稅率 23.0%（含 Anthropic 評價利益的遞延稅）")
inp("debt_q3", "債務面額（2026-09-30，推導：6/30 面額 133.0＋Q3 長期借款收入 48.1 − 償還）", "US$bn", None, fv("debt.face.20260930"), None, "Derived", f"{S_10Q}；{S_PR26Q3}",
    "10-Q 附註未取得；帳面法（長期債帳面變動）為 177.2，差額主要是 2027 年到期票據轉入流動部分")
inp("fl_pv", "融資租賃負債現值（沿用 2026-06-30；10-Q 未取得）", "US$bn", None, fv("lease.finPV.2026q2"), None, "Verified", S_10Q, "附註 4；Q3 新增融資租賃取得 1.47、本金償還 0.43")
inp("fin_obl", "融資義務（financing obligations；沿用 2026-06-30）", "US$bn", None, fv("lease.finObl.2026q2"), None, "Verified", S_10Q, "")
inp("fin_debt", "類債務（融資租賃＋融資義務；淨負債與 WACC 的 D 皆計入，模型期間持平）", "US$bn", None, "=fl_pv+fin_obl", None, "Derived", "修正一（F6、核對第 4 點）", "")
inp("ke", "股權成本＝rf＋β × ERP", "%", None, "=rf+beta*erp", None, "Derived", "CAPM", "")
inp("wacc_adj", "WACC 調整（Sens 用；基準 0）", "pt", -0.01, 0, 0.01, "Assumed", "Sens", "")
inp("wacc", "WACC＝E/(D+E) × ke＋D/(D+E) × kd ×（1 − 稅率）＋調整；E＝評價用股價 × 流通股數；D＝債務面額＋類債務", "%", None,
    "=px_val*sh_out/(px_val*sh_out+debt_q3+fin_debt)*ke+(debt_q3+fin_debt)/(px_val*sh_out+debt_q3+fin_debt)*kd*(1-tax)+wacc_adj", None, "Derived", "修正一（F4）", "")
inp("tg", "永續成長率 g", "%", 0.02, 0.03, 0.04, "Assumed", "工作單第 2 節（g 3%）", "")
inp("ronic", "永續期新投資報酬率（RONIC；預設＝WACC：永續成長不創造價值）", "%", None, "=wacc", "=wacc+0.05", "Assumed", "判斷（v0.3 報告第 5 節；F5）", "高值＝WACC＋5pt（Sens）")
inp("w_dcf", "目標價中 DCF 的權重；其餘為分部加總（SOTP）", "%", 0.3, 0.5, 0.7, "Assumed", "工作單第 2 節（[判斷] 預設 DCF 50%＋分部加總 50%）", "Andy 可改；低／高只用於 Sens")
for s, w in (("A", 0.25), ("B", 0.5), ("C", 0.25)):
    inp(f"w_{s}", f"情境權重 {SCN[s]}", "%", None, w, None, "Assumed", "工作單第 2 節（[判斷] 預設 A 25%／B 50%／C 25%）", "三者合計 100%")
inp("tp_yf", "目標價時點（2027-10-31，約 12 個月後）距 DCF 基準日（2026-12-31）年數", "年", None, "=(DATE(2027,10,31)-DATE(2026,12,31))/365", None, "Derived", "目標價約 12 個月後", "v0.3 為 2027-10-27；DCF 情境價值以 ke 推到目標價時點")
KEYFILL |= {"w_dcf", "w_A", "w_B", "w_C", "tg", "ronic"}

TRN = SNAP["items"]["IF_NonNV_Hold_Trn"]["values"]; TRO = SNAP["items"]["IF_NonNV_Out_Trn"]["values"]
sec("C. 命題模組：AWS AI（對外 AI 在役 MW × 每 MW 年收入）")
inp("mw_ye23", "對外 AI 在役 MW（2023 年底）", "MW", None, 120, None, "Assumed", "沿用 v0.2.2（defaults.mwYearEnd）", "公司未揭露 AI 在役 MW；只用於歷史分拆與折舊年份")
inp("mw_ye24", "對外 AI 在役 MW（2024 年底）", "MW", None, 360, None, "Assumed", "沿用 v0.2.2（defaults.mwYearEnd）", "同上")
inp("mw_ye25", "對外 AI 在役 MW（2025 年底）", "MW", 1100, 1440, 1800, "Assumed", "沿用 v0.2.2（defaults.mwYearEnd）", "Trainium2 1.4M 顆 × 0.871 kW ≈ 1,220 MW-IT（事實總帳 chip.trn2.mwFromLanded）")
inp("mw_open", "對外 AI 在役 MW（2026-06-30；期初機隊錨點）＝在役 AI 2,600 × 對外 80%", "MW", 1400, 2080, 2970, "Derived", f"{S_EPOCH}；{S_V01A}", "Epoch 三站 1,422 MW（Trainium2）＋NVIDIA 推估；第三方估計")
inp("pace_A", "情境 A：對外 AI MW 每年新增（公司「2027 年底容量為 2025 年兩倍」，以 MS 6＋8 GW 對照）", "MW/年", None, 3733, None, "Derived", f"沿用 v0.2.2（scenarios.mwPath.pace.high）；{S_MSTAN}", "每年 +7 GW × AI 占新增 80% ÷ PUE 1.2 × 對外 80%")
inp("pace_B", "情境 B：對外 AI MW 每年新增（C10：2026 資本支出指引約 220 隱含的建置速度；下半年視窗）", "MW/年", None, PACE.get("pace", 3634.3), None, "Derived",
    "amazon/calib_pace_v031.py → data/pace_v031.json（C10 重解）", "（指引 − 上半年實績）÷ 2＝Q4 26 公式資本支出（AI＋非 AI）；把支出換成實體產能，不是以營收反推單價；v0.3 為 3,634.3")
inp("pace_C", "情境 C：對外 AI MW 每年新增（與 B 相同：資本支出照指引投入，下檔在變現）", "MW/年", None, "=pace_B", None, "Assumed", "判斷（v0.3 報告第 5 節）", "")
inp("pace_x", "MW 速度乘數（Sens 用；三情境同乘，評價日前實績不變）", "x", 0.75, 1.0, 1.25, "Assumed", "Sens", "")
inp("mw_q3", "對外 AI 在役 MW（評價日 2026-09-30）＝6/30 機隊＋情境 B 速度 × 1 季", "MW", None, "=mw_open+pace_B*0.25", None, "Derived", "C10 速度（實績期三情境相同）", "公司未揭露；Q3 現金資本支出 75.9")
inp("ext_sh", "對外比例（AI MW 中租給客戶的比例）", "%", 0.6, 0.8, 1.0, "Assumed", "沿用 v0.2.2（對照表 r1 C23：公司未揭露、最大不確定）", "Sens ±20pt")
inp("long_sh", "長約占對外容量比例", "%", 0.7, 0.8, 0.9, "Assumed", "沿用 v0.2.2（pricing.longShare）", "Anthropic Trainium2 三站＋OpenAI 兩筆占多數")
inp("p_spot", "H100 現貨價（Silicon Data 指數）", "US$/GPU-hr", None, 2.82, None, "Verified", S_KEV, "2026-10-06；MAG 三家共用（C1）")
inp("p_list", "AWS p5.48xlarge 隨需（每 GPU 小時）", "US$/GPU-hr", None, 6.88, None, "Verified", S_KEV, "公開牌價（賣價，非成交價）；k 高軸")
inp("p_long", "長約等值價格（判斷倍數 0.75 × v5.27 VR200 成本 12.762）", "US$m/MW-年", None, 9.5715, None, "Assumed", "沿用 v0.2.2（legacy_v0.2/kspec.py long.base）", "判斷值，不是觀察到的市場價格")
inp("p_long_lo", "長約等值價格下緣（判斷倍數 0.55 × 12.762）", "US$m/MW-年", None, 7.0191, None, "Assumed", "沿用 v0.2.2（legacy_v0.2/kspec.py long.low）", "Anthropic 下限 0.22–0.45 只作對照")
for code, nm in (("h100", "H100"), ("gb200", "GB200"), ("gb300", "GB300"), ("vr200", "VR200")):
    inp(f"tk_he_{code}", f"Tokenomics IF_HoldEcon｜{GEN[nm]}（每 MW 年持有成本，經濟口徑）", "US$m/MW-年", tk("IF_HoldEcon", nm, "低成本"), tk("IF_HoldEcon", nm), tk("IF_HoldEcon", nm, "高成本"), "Derived", S_TK,
        "WACC 10% 下的打平持有成本；低／高＝Tokenomics 低成本／高成本欄（Sens 四世代同時，鍵 he_all）")
for code, nm in (("h100", "H100"), ("gb200", "GB200"), ("gb300", "GB300"), ("vr200", "VR200")):
    inp(f"tk_cx_{code}", f"Tokenomics IF_CapexIT｜{GEN[nm]}（每 MW IT 資本支出）", "US$m/MW", None, tk("IF_CapexIT", nm), None, "Derived", S_TK, "")
for code, nm in (("h100", "H100"), ("gb200", "GB200"), ("gb300", "GB300"), ("vr200", "VR200")):
    inp(f"tk_ox_{code}", f"Tokenomics IF_OpexGW｜{GEN[nm]}（每 MW 年營運費用，不含折舊、含電費）", "US$m/MW-年", None, tk("IF_OpexGW", nm), None, "Derived", S_TK, "")
inp("tk_gpuhr", "Tokenomics IF_GPUhrEcon｜Hopper H100（每 GPU 小時持有成本）", "US$/GPU-hr", None, tk("IF_GPUhrEcon", "H100"), None, "Derived", S_TK, "k_現貨的分母")
inp("tk_fac", "Tokenomics L1_FacCapexMW（每 MW 機房資本支出）", "US$m/MW", None, SNAP["items"]["L1_FacCapexMW"]["values"], None, "Derived", S_TK, "")
inp("tk_life", "AI IT 設備折舊年限（Tokenomics IF_DeprLifeIT）", "年", 5, tk("IF_DeprLifeIT", "H100"), 7, "Derived", S_TK, "管理層：伺服器使用 5–6 年（2026Q3 法說）")
inp("fac_life", "機房折舊年限", "年", 15, 20, 30, "Assumed", "沿用 v0.2.2（AI 機房 20 年）", "管理層：資料中心建築 30 年以上")
inp("tk_pq", "Tokenomics L1_RevGW_PQ_VR200（模型商每 MW 年實收 P×Q；只作上限檢查）", "US$m/MW-年", 36.2296, SNAP["items"]["L1_RevGW_PQ_VR200"]["values"], 78.4975, "Derived", S_TK, "上限，不得作預測")
inp("k_long", "k_長約＝長約等值價格 ÷ 目前 VR200 IF_HoldEcon", "x", "=p_long_lo/tk_he_vr200", "=p_long/tk_he_vr200", None, "Derived", "legacy_v0.2/kspec.py（隨目前快照重算）", "低＝長約下緣（Sens）")
inp("k_spot", "k_現貨＝現貨價 ÷ IF_GPUhrEcon｜H100", "x", None, "=p_spot/tk_gpuhr", "=p_list/tk_gpuhr", "Derived", "legacy_v0.2/kspec.py（三家一致，C1）", "高＝AWS 牌價（Sens）")
inp("k_long_lo2", "k_長約下緣＝長約下緣價格 ÷ VR200 IF_HoldEcon（情境 C 用）", "x", None, "=p_long_lo/tk_he_vr200", None, "Derived", "v0.3.1：情境值改用專屬列（原則第 8 節，不讀敏感度欄）", "")
inp("k_spot_lst", "k_現貨（AWS 牌價）＝牌價 ÷ IF_GPUhrEcon｜H100（情境 A 用）", "x", None, "=p_list/tk_gpuhr", None, "Derived", "同上", "")
inp("k_A", "情境 A：新產能 k（價格軸高：現貨改用 AWS 牌價）", "x", None, "=long_sh*k_long+(1-long_sh)*k_spot_lst", None, "Derived", "沿用 v0.2.2 價格軸（pricing.kSpot.high）", "")
inp("k_B", "情境 B：k（價格軸基準：長約 80% × k_長約＋現貨 20% × k_現貨）", "x", None, "=long_sh*k_long+(1-long_sh)*k_spot", None, "Derived", "沿用 v0.2.2（C1 混合 k）", "期初機隊三情境皆用本列")
inp("k_C", "情境 C：新產能 k（價格軸低：長約下緣）", "x", None, "=long_sh*k_long_lo2+(1-long_sh)*k_spot", None, "Derived", "沿用 v0.2.2 價格軸（pricing.kLong.low）", "")
inp("mo_trn2", "期初世代占比：Trainium2（以 Hopper 計價，C2；期初機隊不套 NonNV 係數）", "%", None, 0.55, None, "Assumed", S_V01A, "Epoch 三站 1,422 ÷ 在役 2,600；Tokenomics NonNV 係數只定義 Trainium3 對 VR200")
inp("mo_h100", "期初世代占比：Hopper", "%", None, 0.15, None, "Assumed", S_V01A, "")
inp("mo_gb200", "期初世代占比：GB200", "%", None, 0.15, None, "Assumed", S_V01A, "")
inp("mo_gb300", "期初世代占比：GB300", "%", None, 0.15, None, "Assumed", S_V01A, "")
inp("he_open", "期初機隊加權 IF_HoldEcon", "US$m/MW-年", None, "=(mo_trn2+mo_h100)*tk_he_h100+mo_gb200*tk_he_gb200+mo_gb300*tk_he_gb300", None, "Derived", "上列加權", "")
inp("ox_open", "期初機隊加權營運費用", "US$m/MW-年", None, "=(mo_trn2+mo_h100)*tk_ox_h100+mo_gb200*tk_ox_gb200+mo_gb300*tk_ox_gb300", None, "Derived", "上列加權", "")
inp("rpm_open", "期初機隊每 MW 年收入＝加權 IF_HoldEcon × k_B（收入端不乘 IF_Util，MW-year 計價）", "US$m/MW-年", None, "=he_open*k_B", None, "Derived", "沿用 v0.2.2 口徑", "")
inp("cust_sh", "新增 MW 中自研晶片（Trainium3 → 以 GB300 為錨乘 NonNV 係數）比例", "%", 0.4, 0.5, 0.6, "Assumed", "沿用 v0.2.2（pricing.chips mixAdds）", "Trainium3 2026 年初出貨；2026Q3 法說：Trainium 2027 年底起對外銷售晶片（本版不建模）")
inp("na_vr26", "NVIDIA 新增中 VR200 比例（2026；其餘 GB300）", "%", None, 0, None, "Assumed", "沿用 v0.2.2", "")
inp("na_vr27", "NVIDIA 新增中 VR200 比例（2027）", "%", None, 0.5, None, "Assumed", "沿用 v0.2.2", "")
inp("na_vr28", "NVIDIA 新增中 VR200 比例（2028 起）", "%", None, 1.0, None, "Assumed", "沿用 v0.2.2", "")
inp("trn_hold", "Tokenomics IF_NonNV_Hold_Trn（Trainium3 每 GW 持有成本 ÷ VR200；低／基準／高）", "x", TRN[0], TRN[1], TRN[2], "Assumed", S_TK, "Tokenomics 內為 Assumed；v21.4 為 0.6／0.7／0.9")
inp("trn_out", "Tokenomics IF_NonNV_Out_Trn（Trainium3 每 GW 產出 ÷ VR200；只用於 Sens：收入改按產出比）", "x", TRO[0], TRO[1], TRO[2], "Assumed", S_TK, "Tokenomics PR #172 Q1（收入口徑）未決")
inp("cf_chip", "自研晶片每 MW 收入係數（v0.3.1：改讀 IF_NonNV_Hold_Trn；舊 C2 口徑＝同期 NVIDIA 世代 × 1.0）", "x", None, "=trn_hold", 1.0, "Derived", f"{S_TK}；工作單第 1 節第 3 項（權威在 Tokenomics）", "高＝舊 C2 口徑 1.0（Sens，與 cf_cx 同改）；比照 Alphabet TPU")
inp("cf_cx", "自研晶片每 MW IT 資本支出係數（v0.3.1：改讀 IF_NonNV_Hold_Trn）", "x", None, "=trn_hold", 1.0, "Derived", S_TK, "持有成本比同時作資本支出比（近似：Tokenomics 比值相對 VR200，模型套到 GB300 錨）")
inp("colo_sh", "機房租用比例（第三方託管）", "%", 0.1, 0.2, 0.3, "Interested-party", S_BLOOM, "自建比例＝1 − 本列；租用部分不計機房資本支出")
inp("rent_rec", "租用機房年租金 ÷ 機房資本支出（資本回收係數，C22）", "%", None, 0.10, None, "Assumed", "對照表 r1 C22", "")
inp("rent_mw", "每 MW 年機房租金（攤到全部 AI MW）＝租用比例 × 機房成本 × 回收係數", "US$m/MW-年", None, "=colo_sh*tk_fac*rent_rec", None, "Derived", "上列", "")
inp("capex_g26", "2026 現金資本支出指引（約 220，維持）", "US$bn", None, 220, None, "Interested-party", f"{S_PR26Q3}；{S_CALL3}", "Q4 26 AI 資本支出＝指引 − 9M 實績 − Q4 非 AI（FY2026 錨定指引，修正三）；隱含 Q4 47.8，Q3 為 75.9")
inp("internal_ai_reading", "自用 AI 讀法（2＝讀法 2 主值：自用 MW 以 k＝1 計影子效益、價值中性；1＝讀法 1：只計成本）", "1/2", 1, 2, None, "Assumed", "修正一（F1；對照表 r1 C16；Andy 2026-10-27 15:03 決定讀法 2 為主值）", "讀法 1 列 Sens 與 Val 第 0 節")
inp("cip_lag", "實績期 AI 在建（CIP）轉入在役的延後年數", "年", None, 1, 2, "Assumed", "修正三（獨立核對：AWS 折舊 2024→2025→1H26 年化 13.3→21.5→30.7）", "Sens：2 年")
inp("ai_rr", "公司 AI 營收年化（AWS AI run-rate，2026Q2 >25；Q3 只說年增三位數）", "US$bn", None, fv("aws.aiRunRate.2026q2"), None, "Interested-party", S_PR26Q2, "只作 Recon 驗證列（C3）；不用於校準")
inp("aws_rr", "AWS 年化營收 run-rate（2026Q3）", "US$bn", None, 184, None, "Interested-party", f"{S_PR26Q3}；{S_CALL3}", "只作 Recon 驗證列")
inp("aws_bl", "AWS backlog（2026-09-30）", "US$bn", None, 578, None, "Interested-party", S_CALL3, "年增三位數；只作 Recon 驗證列（不反推 k）")
inp("aws_g40", "公司說法：AWS 2027、2028 年各成長 40% 以上", "%", None, 0.40, None, "Interested-party", f"{S_PR26Q3}；{S_CALL3}", "只作 Recon 驗證列（不校準）")
KEYFILL |= {"pace_A", "pace_B", "pace_C", "ext_sh"}

sec("D. AWS 其他（非 AI 雲端）")
inp("g_awso_lt", "AWS 其他長期年增（2030 起）", "%", 0.05, 0.08, 0.11, "Assumed", "沿用 v0.2.2（gLT 8%）", "2027–2030 由 FY2025 年增（Derived）線性收斂")
inp("awso_g0", "AWS 其他成長起點（0＝FY2025 年增，v0.3 規則；1＝9M26 年增）", "0/1", None, 0, 1, "Assumed", "v0.3 規則（AI 分拆低估時殘差成長偏高）", "高＝9M26 年增（Sens；2026Q3 AWS 年增 39%）")
inp("cx_awso", "AWS 其他資本支出 ÷ AWS 其他營收", "%", 0.15, 0.20, 0.27, "Assumed", "沿用 v0.2.2（cx 20%）", "上緣＝2023 年 AWS PP&E 淨增加 ÷ AWS 營收 27.4%")

sec("E. 零售各營收線（10-K 營收類別；北美＋國際）")
for k, nm, b, lo, hi in (("os", "Online stores", 0.05, 0.03, 0.07), ("ps", "Physical stores", 0.03, 0.01, 0.05), ("sp", "Third-party seller services", 0.07, 0.05, 0.09),
                         ("ad", "Advertising services", 0.10, 0.07, 0.13), ("sub", "Subscription services", 0.07, 0.05, 0.09), ("ot", "Other", 0.05, 0.02, 0.08)):
    inp(f"g_{k}", f"{nm} 長期年增（2030 起）", "%", lo, b, hi, "Assumed", "判斷（v0.3 報告第 5 節）", "Q4 26＝Q4 25 ×（1＋9M26 年增）× 校準係數（合併 Q4 營收＝指引中點）；2027–2030 由 9M26 年增線性收斂到本列")
inp("m_ad", "廣告營業利益率（分攤到北美／國際的方法：見 PL 頁）", "%", 0.40, 0.50, 0.60, "Assumed", "沿用 v0.2.2（廣告 EBITDA 率 50%，分部未揭露）", "數位廣告同業 META、GOOGL EBITDA 率 50%–60%；零售本業利益率＝分部利益扣廣告利潤後的殘差")
inp("step_na", "北美零售本業利益率每年改善（2027–2030）", "pt", 0.0, 0.003, 0.006, "Assumed", "判斷：9M26 對 9M25 改善（Q 頁推導）", "")
inp("step_in", "國際零售本業利益率每年改善（2027–2030）", "pt", -0.003, 0.0, 0.003, "Assumed", "判斷", "")
inp("cx_rt", "零售資本支出 ÷ 零售營收（北美＋國際 PP&E 淨增加，2023–2025 平均）", "%", None, "=AVERAGE(" + ",".join(f"(«Q.pa_na.{c}»+«Q.pa_in.{c}»)/(«Q.na.{c}»+«Q.intl.{c}»)" for c in "CDE") + ")", None, "Derived", S_10K, "4.48%／5.84%／7.40%；含分攤的技術基礎設施，可能偏高（沿用 v0.2.2）")

sec("F. 公司層與現金")
inp("sbc_r", "股份基礎給付 ÷ 營收（近四季 2025Q4–2026Q3）", "%", None, "=«Q.sbc.Q»/«Q.tot.Q»", None, "Derived", S_PR26Q3, "SBC 計入費用；DCF 不加回（視同現金成本）")
inp("oth_r", "營運現金其他項目 ÷ 營收（營運現金流 −（淨利 − 業外淨額 ×（1 − 稅率））− 分部 D&A − SBC；2023–2025 平均；FY2027 起）", "%", -0.005, "=AVERAGE(" + ",".join(f"(«Q.cfo.{c}»-(«Q.ni.{c}»-«Q.nonop.{c}»*(1-«Q.etr.{c}»))-«Q.da_seg.{c}»-«Q.sbc.{c}»)/«Q.tot.{c}»" for c in "CDE") + ")", 0.02, "Derived", S_10K,
    "0.6%／0.7%／2.3%；扣除業外（持股評價利益為非現金）")
inp("oth_q4", "Q4 26 其他營運項目 ÷ 營收（2025Q4 同期實績；C14：首期營運資金用前一年度同期）", "%", None, "=(«Q.cfo.J»-(«Q.ni.J»-«Q.nonop.J»*(1-«Q.etr.E»))-«Q.da_seg.J»-«Q.sbc.J»)/«Q.tot.J»", None, "Derived", S_PR25Q4,
    "Amazon 第四季營運資金季節性流入（應付帳款）；v0.3 以全年平均 oth_r 計 2H")
inp("lease_pr", "融資租賃本金償還（每年）", "US$bn", 1.5, 1.6, 3.0, "Assumed", S_10K, "2023 4.384、2024 2.043、2025 1.557；9M26 1.296")
inp("sh_g", "稀釋股數年增（無回購，SBC 稀釋）", "%", 0.0, 0.008, 0.015, "Derived", S_XBRL, "稀釋加權股數 9M25→9M26 +0.8%")
inp("kd_book", "既有債務平均利率（2026Q3 利息費用 × 4 ÷ 平均債務面額）", "%", None, "=«Q.intx.M»*4/((«Q.debt.L»+«Q.debt.M»)/2)", None, "Derived", S_PR26Q3, "既有債務以本列計息；新增借款以 kd_new")
inp("y_cash", "現金與有價證券收益率（2026Q3 利息收入 × 4 ÷ 期末現金）", "%", None, "=«Q.inti.M»*4/«Q.cash.M»", None, "Derived", S_PR26Q3, "")
inp("cash_q3", "現金、約當現金與有價證券（2026-09-30）", "US$bn", None, fv("bs.cashMs.2026q3"), None, "Interested-party", S_PR26Q3, "55.8＋58.4（v0.3 6/30 為 123.0）")
inp("min_cash", "最低現金（低於即舉債）", "US$bn", 30, 50, 80, "Assumed", "沿用 v0.2.2（minCash 50）", "")
inp("inv_q4", "2026Q4 策略投資", "US$bn", None, 0, 15.0, "Assumed", f"{S_10Q}；{S_PR26Q3}", "Q3 收購與投資 24.0（含 OpenAI 承諾的期後補足 21.3）；Anthropic 可動用額度 15.0 依里程碑，高＝全數於 Q4 動用")
inp("dps", "每股股利", "US$", None, 0, None, "Verified", S_PR26Q3, "Amazon 未發放現金股利；9M26 無回購（現金流量表融資段無買回列）")

sec("G. 持股（分部加總與 DCF 都加；預測期業外評價損益設 0）")
inp("ant_val", "Anthropic 投後估值（Series H）", "US$bn", None, fv("rel.ant.valSeriesH"), None, "Interested-party", S_ANT, "新創自報估值；2026Q3 其他業外 −0.7（無新評價利益）")
inp("ant_sh", "Amazon 持有 Anthropic 比例", "%", 0.15, 0.18, 0.21, "Assumed", f"沿用 v0.2.2；{S_FORT}", "Fortune 15%–19%、帳列 ÷ 投後 19.7%")
inp("oai_val", "OpenAI 投後估值（投前 730＋該輪 110）", "US$bn", None, 840, None, "Derived", S_OAI, "")
inp("oai_sh", "Amazon 持有 OpenAI 比例（投入 50 ÷ 840）", "%", None, 0.0595, None, "Derived", "沿用 v0.2.2", "資本結構（PBC＋基金會）可能使比例不同")
inp("liq_disc", "持股流動性折價", "%", 0.0, 0.2, 0.4, "Assumed", "工作單第 2 節（20%）", "")
inp("m_ai", "AWS AI 的 EV/EBITDA 倍數（只在 ai_leg＝0 時使用）", "x", 4, 6, 10, "Assumed", "沿用 v0.2.2（對照表 D6）", "CRWV TTM 25.0×、ORCL 16.5×（只作對照）")
inp("sbc_ab", "分部加總 EBITDA 是否加回 SBC（1＝加回，與同業倍數同口徑；0＝視同成本）", "0/1", 0, 1, None, "Assumed", "修正二（F2 敏感度）", "")
inp("ai_leg", "分部加總中 AWS AI 的評價法（1＝自身 DCF，主值；0＝EBITDA × m_ai 倍數，對照）", "0/1", 0, 1, None, "Assumed", "修正一（F3）", "AI 是資本密集業務，EBITDA 倍數忽略汰換與成長資本支出")
inp("m_x", "分部倍數乘數（Sens 用；基準 1）", "x", 0.8, 1.0, 1.2, "Assumed", "Sens", "")

sec("H. 市場共識與公司指引（只作 Recon 對照，不校準；Q4 營收指引中點用於零售 Q4 校準）")
AE = CONS["annualEstimates"]
for y in (2026, 2027, 2028):
    e = AE[f"FY{y}"]
    for k, lab, v in (("rev", "營收", e["revenue"]), ("ebitda", "EBITDA", e["ebitda"]), ("ebit", "EBIT", e["ebit"]), ("ni", "淨利（GAAP）", e["netIncome"]),
                      ("eps", "EPS（GAAP）", e["epsGaap"]), ("capex", "資本支出", e["capex"]), ("fcf", "自由現金流", e["fcf"])):
        inp(f"c_{k}{y % 100}", f"共識 FY{y} {lab}（MarketScreener，財報前）", "US$" if k == "eps" else "US$bn", None, v, None, "Interested-party", S_MS, "2026-10-08 擷取；財報後版本未取得")
inp("c_sprev26", "共識 FY2026 營收（S&P Global，2026-10-26，財報前）", "US$bn", 815.9, 828.74, 842.0, "Interested-party", S_SP, "56 位；" + NM)
inp("c_sprev27", "共識 FY2027 營收（S&P Global，2026-10-26）", "US$bn", None, 949.92, None, "Interested-party", S_SP, "")
inp("c_speps26", "共識 FY2026 調整後 EPS（S&P Global，2026-10-26）", "US$", 7.65, 7.92, 8.30, "Interested-party", S_SP, "non-GAAP（不含 Anthropic 評價利益）；" + NM)
inp("c_speps27", "共識 FY2027 調整後 EPS（S&P Global，2026-10-26）", "US$", None, 10.43, None, "Interested-party", S_SP, "")
inp("cons_mb", "市場共識目標價（MarketBeat 60 位，2026-10-29）", "US$", None, 325.12, None, "Interested-party", S_MB, "第二個獨立彙整（底層為各券商報告）")
inp("g26q4_lo", "公司指引 2026Q4 營收下緣", "US$bn", None, 240.0, None, "Interested-party", S_PR26Q3, "年增 12%–16%；匯率約 −60bp")
inp("g26q4_hi", "公司指引 2026Q4 營收上緣", "US$bn", None, 247.0, None, "Interested-party", S_PR26Q3, "")
inp("g26q4_oilo", "公司指引 2026Q4 營業利益下緣", "US$bn", None, 28.5, None, "Interested-party", S_PR26Q3, "Leo 成本 Q4 多數仍費用化（法說）")
inp("g26q4_oihi", "公司指引 2026Q4 營業利益上緣", "US$bn", None, 33.5, None, "Interested-party", S_PR26Q3, "")
inp("g26q4_mid", "公司指引 2026Q4 營收中點（Q4 零售營收校準目標）", "US$bn", None, "=(g26q4_lo+g26q4_hi)/2", None, "Derived", S_PR26Q3, "")

sec("I. v0.3 對照（Recon 變動拆解用）")
inp("v03_tp", "v0.3 加權目標價（讀法 2；2026-10-27）", "US$", None, 215.58, None, "Derived", S_V03, "215.578882（v0.3 成品 Val!B19）")
inp("v03_tp1", "v0.3 讀法 1 加權目標價", "US$", None, 194.26, None, "Derived", S_V03, "")
inp("v03_px", "v0.3 現價（2026-10-26）", "US$", None, 249.09, None, "Verified", S_V03, "")

INROW = {}
r = 4
for row in ROWS:
    if row[0] == "sec":
        IN.cell(r, 1, row[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = row
    if key in OVR:
        base = OVR[key]
    if key == "v03_tp":
        base = 215.578882385843
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            c = IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")):
                c.font = BLUE
            c.number_format = "0.0%" if unit in ("%", "pt") else ("#,##0.0000" if unit in ("x",) else "#,##0.00##")
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}")
    if lo is not None: defname(key + "_lo", f"In!$D${r}")
    if hi is not None: defname(key + "_hi", f"In!$F${r}")
    if key in KEYFILL:
        for c in range(1, 10): IN.cell(r, c).fill = KEY
    INROW[key] = r; r += 1
# v0.3.1：J. 情境驅動對照（原則第 8 節；純展示，只引用上方各列；三情境並排，與上方「敏感度低／高」欄分開）
r += 1
IN.cell(r, 1, "J. 情境驅動對照（三情境並排，進主模型的加權目標價；本節只引用上方各列，修改請改上方藍字。上方「敏感度低／高」欄只用於 Sens 頁，不進目標價）").font = BOLD; r += 1
hdr(IN, r, ["鍵", "項目", "單位", "情境 A（樂觀）", "情境 B（基準）", "情境 C（保守）", "標記", "說明"]); r += 1
J0 = r
for key, item, unit, ks, fmt, note in (
        ("scn_pace", "對外 AI MW 每年新增", "MW/年", ("pace_A", "pace_B", "pace_C"), "#,##0.0", "A＝公司容量目標速度；B＝C10；C＝與 B 相同（資本支出照指引投入）"),
        ("scn_k", "新產能 k（每 MW 年收入 ÷ 持有成本）", "x", ("k_A", "k_B", "k_C"), "0.000", "A＝價格軸高；B＝基準；C＝價格軸低；期初機隊三情境皆用 k_B"),
        ("scn_w", "情境權重", "%", ("w_A", "w_B", "w_C"), "0%", "[判斷] 黃底，Andy 可改")):
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, k_ in zip((4, 5, 6), ks):
        IN.cell(r, j, f"=E{INROW[k_]}").number_format = fmt
    IN.cell(r, 7, "Derived"); IN.cell(r, 8, "引用 In 頁 " + "／".join(ks) + "；" + note)
    r += 1
JROWS = list(range(J0, r))
for col, w in zip("ABCDEFGHI", [12, 70, 12, 10, 12, 10, 16, 60, 70]):
    IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# =====================================================================================
# 參照記號：«區塊.鍵.欄» → 工作表!$欄$列（所有區塊寫完後一次解析，允許前後互相引用）
# =====================================================================================
REG = {}
TOK = re.compile(r"«([^.»]+)\.([^.»]+)\.([A-Z]+)»")


def resolve_all():
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "«" in c.value:
                    def rep(m):
                        k = (m.group(1), m.group(2))
                        if k not in REG:
                            raise KeyError(f"未定義的參照 {m.group(0)}（{ws.title}!{c.coordinate}）")
                        sh, rr = REG[k]
                        return f"{sh}!${m.group(3)}${rr}"
                    c.value = TOK.sub(rep, c.value)
                    c.value = RNG2.sub(r"\1!\2:", c.value)


RNG2 = re.compile(r"([A-Za-z]+)!(\$[A-Z]+\$\d+):\1!")

# =====================================================================================
# Q：最近 8 季與最近完整會計年度（facts，SEC 一手資料）
# =====================================================================================
Q = wb.create_sheet("Q")
Q["A1"] = ("Q — 實績（US$bn；股數十億股）：FY2023–FY2025、最近 8 季（2024Q4–2026Q3）、9M25、9M26、1H26 與近四季。資料取自 data/amazon_facts_20261008.json（10-K、10-Q、財報新聞稿逐季補充表）"
           "與 data/amazon_facts_2026q3.json（2026Q3 新聞稿，兩讀核對；10-Q 未取得），藍字＝事實值、黑字＝加總或換算")
Q["A2"] = "核對列（綠底）：零售營收線合計＝北美＋國際、營收線合計＝合併營收、分部合計＝合併營收、分部營業利益合計＋未分攤＝合併營業利益、9M26＝事實總帳九個月值；實績皆須為 0"
QCOL = {"fy2023": "C", "fy2024": "D", "fy2025": "E", "2024q4": "F", "2025q1": "G", "2025q2": "H", "2025q3": "I",
        "2025q4": "J", "2026q1": "K", "2026q2": "L", "2026q3": "M"}
hdr(Q, 3, ["項目", "單位", "FY2023A", "FY2024A", "FY2025A", "2024Q4", "2025Q1", "2025Q2", "2025Q3", "2025Q4", "2026Q1", "2026Q2", "2026Q3",
           "9M25", "9M26", "1H26", "近四季（2025Q4–2026Q3）", "標記", "來源"])
QR = {}
qr = [4]
FLOWSUM = (("N", "GHI"), ("O", "KLM"), ("P", "KL"), ("Q", "JKLM"))


def qsec(t):
    Q.cell(qr[0], 1, t).font = BOLD; qr[0] += 1


def qrow(key, label, idbase=None, unit="US$bn", flow=True, scale=1000.0, cells=None, tag="Verified", srcs="", fmt="#,##0.000", check=False):
    r = qr[0]
    Q.cell(r, 1, label); Q.cell(r, 2, unit)
    if idbase:
        for per, col in QCOL.items():
            fid = f"{idbase}.{per}"
            if fid in FACTS and isinstance(FACTS[fid].get("value"), (int, float)):
                c = Q[f"{col}{r}"]; c.value = FACTS[fid]["value"] / scale; c.font = BLUE
    for col, f in (cells or {}).items():
        Q[f"{col}{r}"] = f
        Q[f"{col}{r}"].font = Font(color="000000") if (isinstance(f, str) and f.startswith("=")) else BLUE
    if flow:
        for col, parts in FLOWSUM:
            if Q[f"{col}{r}"].value is None and all(Q[f"{p}{r}"].value is not None for p in parts):
                Q[f"{col}{r}"] = "=" + "+".join(f"{p}{r}" for p in parts)
    for col in "CDEFGHIJKLMNOPQ":
        Q[f"{col}{r}"].number_format = fmt
    Q.cell(r, 18, tag); Q.cell(r, 19, srcs)
    if check:
        for col in range(1, 18): Q.cell(r, col).fill = CHK
    QR[key] = r; REG[("Q", key)] = ("Q", r); qr[0] += 1
    return r


SEGSRC = f"{S_10K}（FY）；{S_PR26Q3}、{S_PR26Q2}、{S_PR24Q4}（逐季補充表）"
qsec("1. 官方報告分部營收（10-K 附註：North America、International、AWS）")
qrow("na", "North America", "is.NA.rev", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("intl", "International", "is.INTL.rev", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("aws", "AWS", "is.AWS.rev", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("tot", "合併營收", "is.TOT.rev", tag="Verified／Interested-party", srcs=SEGSRC)
qsec("2. 營收線（10-K 營收類別；零售各線＝北美＋國際，AWS 線＝AWS 分部）")
LINES = (("os", "Online stores", "is.cat.online"), ("ps", "Physical stores", "is.cat.physical"), ("sp", "Third-party seller services", "is.cat.seller3p"),
         ("ad", "Advertising services", "is.cat.ads"), ("sub", "Subscription services", "is.cat.subs"), ("ot", "Other", "is.cat.other"))
for k, lab, idb in LINES:
    qrow(k, lab, idb, tag="Verified／Interested-party", srcs=SEGSRC)
ALLQ = "CDEFGHIJKLMNOPQ"
qrow("rt", "零售營收線合計（不含 AWS）", cells={c: "=" + "+".join(f"{c}{QR[k]}" for k, _, _ in LINES) for c in ALLQ}, flow=False, tag="Derived")
qrow("awsl", "AWS（營收線）", cells={c: f"={c}{QR['aws']}" for c in ALLQ}, flow=False, tag="Derived")
qrow("chk_rt", "檢查：零售營收線合計 − (North America＋International)", cells={c: f"={c}{QR['rt']}-({c}{QR['na']}+{c}{QR['intl']})" for c in ALLQ}, flow=False, tag="Derived", check=True)
qrow("chk_ln", "檢查：營收線合計（含 AWS）− 合併營收", cells={c: f"={c}{QR['rt']}+{c}{QR['awsl']}-{c}{QR['tot']}" for c in ALLQ}, flow=False, tag="Derived", check=True)
qrow("chk_seg", "檢查：分部合計 − 合併營收", cells={c: f"={c}{QR['na']}+{c}{QR['intl']}+{c}{QR['aws']}-{c}{QR['tot']}" for c in ALLQ}, flow=False, tag="Derived", check=True)
qrow("chk_9m", "檢查：9M26 合併營收 − 新聞稿九個月值（事實 is.TOT.rev.2026m9）", cells={"O": f"=O{QR['tot']}-{fv('is.TOT.rev.2026m9') if 'is.TOT.rev.2026m9' in FACTS else 0}"}, flow=False, tag="Derived", check=True)
qsec("3. 分部營業利益")
qrow("oi_na", "North America 營業利益", "is.NA.oi", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("oi_in", "International 營業利益", "is.INTL.oi", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("oi_aws", "AWS 營業利益", "is.AWS.oi", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("oi", "合併營業利益", "is.TOT.oi", tag="Verified／Interested-party", srcs=SEGSRC)
qrow("unal", "未分攤（合併 − 分部合計；Amazon 分部已含全部費用）", cells={c: f"={c}{QR['oi']}-({c}{QR['oi_na']}+{c}{QR['oi_in']}+{c}{QR['oi_aws']})" for c in ALLQ}, flow=False, tag="Derived")
qrow("chk_oi", "檢查：分部營業利益合計＋未分攤 − 合併營業利益", cells={c: f"={c}{QR['oi_na']}+{c}{QR['oi_in']}+{c}{QR['oi_aws']}+{c}{QR['unal']}-{c}{QR['oi']}" for c in ALLQ}, flow=False, tag="Derived", check=True)
qrow("chk_9moi", "檢查：9M26 合併營業利益 − 新聞稿九個月值（事實 is.TOT.oi.2026m9）", cells={"O": f"=O{QR['oi']}-{fv('is.TOT.oi.2026m9') if 'is.TOT.oi.2026m9' in FACTS else 0}"}, flow=False, tag="Derived", check=True)
qsec("4. 分部折舊攤銷（不含影音內容與營業租賃資產攤銷；2024 季度未收；2025Q4＝全年 − 前三季；2026Q3＝2026Q2 × 現金流量表 D&A 季增比〔推導，10-Q 未取得〕）")
for k, idb, lab in (("da_na", "is.NA.da", "North America 折舊攤銷"), ("da_in", "is.INTL.da", "International 折舊攤銷"), ("da_aws", "is.AWS.da", "AWS 折舊攤銷")):
    r = qrow(k, lab, idb, flow=False, tag="Verified／Derived（2026Q3）", srcs=f"{S_10K}；{S_10Q}；{S_10Q3}")
    Q[f"J{r}"] = f"=E{r}-G{r}-H{r}-I{r}"; Q[f"J{r}"].font = Font(color="000000")
    Q[f"M{r}"] = f"=L{r}*M{{DACF}}/L{{DACF}}"
    for col, parts in FLOWSUM:
        Q[f"{col}{r}"] = "=" + "+".join(f"{p}{r}" for p in parts)
qrow("da_seg", "分部折舊攤銷合計", cells={c: f"=IF(COUNT({c}{QR['da_na']}:{c}{QR['da_aws']})=3,{c}{QR['da_na']}+{c}{QR['da_in']}+{c}{QR['da_aws']},\"\")" for c in ALLQ}, flow=False, tag="Derived")
qsec("5. 損益（業外、稅、淨利）")
qrow("intx", "利息費用", "is.intExp", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
qrow("inti", "利息收入", "is.intInc", flow=False, tag="Interested-party", srcs=f"{S_10Q}；{S_PR26Q3}")
qrow("pretax", "稅前淨利", "is.pretax", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
qrow("taxx", "所得稅", "is.tax", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
qrow("ni", "淨利", "is.ni", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
qrow("nonop", "業外淨額（稅前 − 營業利益＋利息費用；含 Anthropic 等持股評價利益）",
     cells={c: f"=IF(AND(ISNUMBER({c}{QR['pretax']}),ISNUMBER({c}{QR['intx']})),{c}{QR['pretax']}-{c}{QR['oi']}+{c}{QR['intx']},\"\")" for c in ALLQ}, flow=False, tag="Derived")
qrow("etr", "有效稅率", cells={c: f"={c}{QR['taxx']}/{c}{QR['pretax']}" for c in "CDEO"}, flow=False, tag="Derived", fmt="0.00%")
_shd = {"N": fv("sh.wDiluted.2025m9"), "O": fv("sh.wDiluted.2026m9")}
qrow("shd", "稀釋加權股數", "sh.wDiluted", unit="十億股", flow=False, tag="Verified／Interested-party", srcs=f"{S_XBRL}；{S_PR26Q3}", fmt="#,##0.000", cells=_shd)
qrow("eps", "稀釋 EPS（淨利 ÷ 稀釋加權股數）", cells={c: f"=IF(ISNUMBER({c}{QR['shd']}),{c}{QR['ni']}/{c}{QR['shd']},\"\")" for c in "CDEFGHIJKLMNO"}, unit="US$", flow=False, tag="Derived", fmt="0.00")
qsec("6. 現金流量、資本支出與股東回饋")
qrow("cfo", "營運現金流", "cf.cfo", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}；{S_PR26Q2}；{S_PR26Q1}；{S_PR25Q4}；{S_PR25Q3}")
qrow("cxg", "資本支出（購置 PP&E，總額）", "cf.capexGross", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
qrow("cxn", "資本支出（現金，淨額＝購置 − 出售及獎勵）", "cf.capexNet", tag="Derived", srcs="由同季新聞稿現金流量表換算")
r_dacf = qrow("da_cf", "折舊攤銷（現金流量表；含影音內容與營業租賃資產）", "cf.da", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
for k in ("da_na", "da_in", "da_aws"):
    Q[f"M{QR[k]}"] = Q[f"M{QR[k]}"].value.replace("{DACF}", str(r_dacf)); Q[f"M{QR[k]}"].font = Font(color="000000")
qrow("sbc", "股份基礎給付", "cf.sbc", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}")
qrow("fcf", "自由現金流（公司口徑＝營運現金流 − 資本支出淨額）", cells={c: f"={c}{QR['cfo']}-{c}{QR['cxn']}" for c in ALLQ}, flow=False, tag="Derived")
qrow("lpr", "融資租賃本金償還", "cf.finLeasePrin", tag="Verified／Interested-party", srcs=f"{S_10K}；{S_PR26Q3}", flow=False,
     cells={"N": fv("cf.finLeasePrin.2025m9"), "O": fv("cf.finLeasePrin.2026m9")})
qrow("bb", "買回庫藏股", "cf.buyback", flow=False, tag="Verified", srcs=f"{S_10K}；{S_PR26Q3}", cells={"O": 0})
qrow("div", "現金股利", "cf.dividend", flow=False, tag="Verified", srcs=f"{S_10K}；{S_PR26Q3}", cells={"O": 0})
qrow("acq", "策略投資（收購與非上市投資淨額；2026 含 OpenAI、Anthropic）", flow=False, tag="Interested-party", srcs=f"{S_10Q}；{S_PR26Q3}",
     cells={"P": fv("cf.acqInvest.2026h1"), "O": fv("cf.acqInvest.2026m9")})
qsec("7. 資本支出分布（PP&E 淨增加，含融資租賃取得與已取得未付款；10-K／10-Q 分部附註；2026Q3 未取得）")
for k, idb, lab in (("pa_na", "capex.netAdd.NA", "North America PP&E 淨增加"), ("pa_in", "capex.netAdd.INTL", "International PP&E 淨增加"),
                    ("pa_aws", "capex.netAdd.AWS", "AWS PP&E 淨增加"), ("pa_corp", "capex.netAdd.CORP", "公司層 PP&E 淨增加"), ("pa_tot", "capex.netAdd.TOT", "PP&E 淨增加合計")):
    r = qrow(k, lab, idb, flow=False, tag="Verified", srcs=f"{S_10K}；{S_10Q}")
    Q[f"P{r}"] = fv(f"{idb}.2026h1"); Q[f"P{r}"].font = BLUE
    Q[f"L{r}"] = None
qsec("8. 期末餘額")
qrow("cash", "現金、約當現金與有價證券", flow=False, tag="Verified／Interested-party", srcs=f"{S_10K}；{S_10Q}；{S_PR26Q3}",
     cells={"E": fv("bs.cash.fy2025") + fv("bs.mktSec.fy2025"), "L": fv("bs.cashMs.2026q2"), "M": fv("bs.cashMs.2026q3")})
qrow("debt", "債務面額（2026-09-30 為推導值：6/30 面額＋Q3 長期借款收入 − 償還；10-Q 未取得）", flow=False, tag="Verified／Derived", srcs=f"{S_10K}；{S_10Q}；{S_PR26Q3}",
     cells={"E": fv("debt.face.fy2025"), "L": fv("debt.face.2026q2"), "M": fv("debt.face.20260930")})
for col in range(3, 18):
    Q.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 11
Q.column_dimensions["A"].width = 60; Q.column_dimensions["R"].width = 22; Q.column_dimensions["S"].width = 60
Q.freeze_panes = "C4"
# =====================================================================================
# 引擎：一個情境一組區塊（Thesis → Seg → PL → CF → Val）；Sens 頁以同一函式產生案例區塊
# 欄位：C FY2024A｜D FY2025A｜E 9M26A｜F Q4 26E｜G FY2026E（＝9M 實績＋Q4 模型）｜H–M FY2027–FY2032｜N 9M25A｜O Q4 25A｜P 說明｜R–AA 年份輔助欄（2023–2032）
# =====================================================================================
HIST = ["C", "D", "N", "O", "E"]
QM = {"C": "D", "D": "E", "N": "N", "O": "J", "E": "O"}          # 引擎實績欄 → Q 頁欄（9M25、Q4 25、9M26）
FUT = list(zip(range(2027, 2033), "HIJKLM"))
FC = [c for _, c in FUT]
PREV = {"H": "G", "I": "H", "J": "I", "K": "J", "L": "K", "M": "L"}
YEAR = {"C": 2024, "D": 2025, "G": 2026, **{c: y for y, c in FUT}}
VC = {y: openpyxl.utils.get_column_letter(18 + i) for i, y in enumerate(range(2023, 2033))}   # R..AA
VY = {2024: "C", 2025: "D", 2026: "G", **{y: c for y, c in FUT}}
ALL = ["C", "D", "N", "O", "E", "F", "G"] + FC
HEAD = ["項目", "單位", "FY2024A", "FY2025A", "9M26A", "Q4 26E", "FY2026E（9M 實績＋Q4 模型）"] + [f"FY{y}E" for y, _ in FUT] + ["9M25A", "Q4 25A", "說明", None] + list(range(2023, 2033))
CUR = {}


class Blk:
    def __init__(self, bid, ws, title, head=True):
        self.id, self.ws = bid, ws
        CUR.setdefault(ws.title, 4)
        ws.cell(self.r, 1, title).font = Font(bold=True, size=12); self.r += 1
        if head:
            hdr(ws, self.r, HEAD)
            REG[(bid, "hdr")] = (ws.title, self.r); self.r += 1

    @property
    def r(self):
        return CUR[self.ws.title]

    @r.setter
    def r(self, v):
        CUR[self.ws.title] = v

    def t(self, key, col):
        return f"«{self.id}.{key}.{col}»"

    def sec(self, text):
        self.ws.cell(self.r, 1, text).font = BOLD; self.r += 1

    def row(self, key, label, unit, cells, fmt="#,##0.00", note="", check=False, bold=False, key_fill=False):
        r = self.r
        self.ws.cell(r, 1, label); self.ws.cell(r, 2, unit)
        for col, f in cells.items():
            c = self.ws[f"{col}{r}"]; c.value = f; c.number_format = fmt
            if not (isinstance(f, str) and f.startswith("=")):
                c.font = BLUE
        if note:
            self.ws.cell(r, 16, note)
        if bold:
            self.ws.cell(r, 1).font = BOLD
        if check:
            for col in range(1, 16): self.ws.cell(r, col).fill = CHK
        if key_fill:
            for col in range(1, 16): self.ws.cell(r, col).fill = KEY
        REG[(self.id, key)] = (self.ws.title, r); self.r += 1
        return r

    def close(self):
        self.r = self.r + 2


def yrs(b):
    rr = REG[(b.id, "hdr")][1]
    return f"{b.ws.title}!$R${rr}:$AA${rr}"


def fade(b, key, glt, c0="E"):
    """年增：2027–2030 由 c0 欄（近期年增）線性收斂到 glt；2031–2032＝glt。"""
    out = {}
    for n, (y, c) in enumerate(FUT, 1):
        out[c] = f"={b.t(key, c0)}+({glt}-{b.t(key, c0)})*{min(n, 4)}/4" if n <= 4 else f"={glt}"
    return out


def engine(eid, s, sheets, over=None, tag=""):
    over = over or {}
    dflt = {"pace": f"(pace_{s}*pace_x)", "kn": f"k_{s}"}

    def P(k):
        return over.get(k, dflt.get(k, k))
    nm = f"情境 {SCN[s]}{tag}"

    # ---------------- Seg（一）：零售各營收線與地區分攤 ----------------
    S = Blk(f"{eid}S", sheets["Seg"], f"Seg｜{nm}：營收分部層（官方分部 North America、International、AWS；10-K 營收線）")
    S.sec("1. 零售各營收線（Q4 26＝Q4 25 ×（1＋9M26 年增）× 校準係數；2027–2030 年增由 9M26 年增線性收斂到長期值；2031–2032＝長期值）")
    for k, lab, _ in LINES:
        cells = {c: f"=«Q.{k}.{QM[c]}»" for c in HIST}
        cells["F"] = f"={S.t('rev_' + k, 'O')}*(1+{S.t('g_' + k, 'F')})*{S.t('cal', 'F')}"
        cells["G"] = f"={S.t('rev_' + k, 'E')}+{S.t('rev_' + k, 'F')}"
        for y, c in FUT:
            cells[c] = f"={S.t('rev_' + k, PREV[c])}*(1+{S.t('g_' + k, c)})"
        S.row("rev_" + k, lab, "US$bn", cells, note="Q 頁實績；預測＝前期 ×（1＋年增）")
        g = {"E": f"={S.t('rev_' + k, 'E')}/{S.t('rev_' + k, 'N')}-1", "F": f"={S.t('g_' + k, 'E')}", "G": f"={S.t('rev_' + k, 'G')}/{S.t('rev_' + k, 'D')}-1"}
        g.update(fade(S, "g_" + k, P(f"g_{k}"), "F"))
        S.row("g_" + k, f"  {lab} 年增", "%", g, "0.0%", note=f"E＝9M26 對 9M25（Prime Day 移入 Q2 的時點影響在九個月內抵銷）；F＝Q4 26 成長基準＝E（Derived）；收斂到 In 頁 g_{k}")
    S.row("cal", "Q4 26 零售營收校準係數（使情境 B 的 Q4 26 合併營收＝指引中點）", "x",
          {"F": f"=(g26q4_mid-«BT.rev.F»-«BT.awso.F»)/(" + "+".join(f"{S.t('rev_' + k, 'O')}*(1+{S.t('g_' + k, 'F')})" for k, _, _ in LINES) + ")"},
          "0.0000", note="Q4 26 合併營收目標＝指引中點 243.5 − 主頁情境 B 的 AWS Q4 26；三情境共用同一係數（修正二），零售不隨 AI 分拆變動")
    S.row("rt", "零售營收線合計（＝North America＋International）", "US$bn", {c: "=" + "+".join(S.t("rev_" + k, c) for k, _, _ in LINES) for c in ALL}, bold=True)
    sh = {c: f"=«Q.na.{QM[c]}»/(«Q.na.{QM[c]}»+«Q.intl.{QM[c]}»)" for c in HIST}
    sh.update({"F": f"={S.t('na_sh', 'O')}", "G": f"={S.t('na', 'G')}/{S.t('rt', 'G')}", **{c: f"={S.t('na_sh', 'G')}" for c in FC}})
    S.row("na_sh", "北美占零售營收比例（地區分攤）", "%", sh, "0.0%", note="Q4 26＝Q4 25 比例；2027 起固定為 FY2026（Derived）；各營收線不分地區揭露，故以本列分攤")
    na = {c: f"=«Q.na.{QM[c]}»" for c in HIST}
    na.update({"F": f"={S.t('rt', 'F')}*{S.t('na_sh', 'F')}", "G": f"={S.t('na', 'E')}+{S.t('na', 'F')}", **{c: f"={S.t('rt', c)}*{S.t('na_sh', c)}" for c in FC}})
    S.row("na", "North America", "US$bn", na, bold=True, note="官方分部；預測＝零售營收線合計 × 北美比例")
    it = {c: f"=«Q.intl.{QM[c]}»" for c in HIST}
    it.update({"F": f"={S.t('rt', 'F')}*(1-{S.t('na_sh', 'F')})", "G": f"={S.t('intl', 'E')}+{S.t('intl', 'F')}", **{c: f"={S.t('rt', c)}*(1-{S.t('na_sh', c)})" for c in FC}})
    S.row("intl", "International", "US$bn", it, bold=True, note="官方分部；預測＝零售營收線合計 ×（1 − 北美比例）")
    S.row("ad_na", "廣告營收歸北美（＝廣告 × 北美比例）", "US$bn", {c: f"={S.t('rev_ad', c)}*{S.t('na_sh', c)}" for c in ALL}, note="廣告營收依零售營收的地區比例分攤（Assumed；公司未分地區揭露）")
    S.row("ad_in", "廣告營收歸國際", "US$bn", {c: f"={S.t('rev_ad', c)}-{S.t('ad_na', c)}" for c in ALL})

    # ---------------- Thesis：AWS AI ----------------
    T = Blk(f"{eid}T", sheets["Thesis"], f"Thesis｜{nm}：命題模組 AWS AI＝對外 AI 在役 MW × 每 MW 年收入（Tokenomics 持有成本 × k）；AI 資本支出 → 折舊 → 分部營業利益")
    t = T.t
    T.sec("1. 容量（對外 AI 在役 MW-IT；期末）")
    mw = {"C": "=mw_ye24", "D": "=mw_ye25", "N": "=mw_ye24+(mw_ye25-mw_ye24)*0.75", "O": "=mw_ye25", "E": "=mw_q3", "F": f"={t('mw', 'E')}+{P('pace')}*0.25", "G": f"={t('mw', 'F')}"}
    mw.update({c: f"={t('mw', PREV[c])}+{t('pace_y', c)}" for c in FC})
    py = {c: f"={P('pace')}" for c in "HIJ"}
    py.update({c: f"={P('pace')}+(«{eid}T.mw.J»*{P('tg')}-{P('pace')})*{n}/3" for n, c in enumerate("KLM", 1)})
    T.row("pace_y", "對外 AI MW 每年新增（2030–2032 線性收斂到 FY2029 年底 MW × g，F5）", "MW/年", py, "#,##0", note="終值不落在建置高峰（共同修正 F5）")
    T.row("mw", "對外 AI 在役 MW（期末）", "MW", mw, "#,##0", note="評價日 2026-09-30＝In 頁 mw_q3（6/30 機隊＋C10 速度一季）；之後每年新增＝In 頁 pace（情境）× pace_x；歷史年底值只用於分拆與折舊年份", key_fill=True)
    avg = {"C": "=(mw_ye23+mw_ye24)/2", "D": f"=({t('mw', 'C')}+{t('mw', 'D')})/2", "N": f"=({t('mw', 'C')}+{t('mw', 'N')})/2", "O": f"=({t('mw', 'N')}+{t('mw', 'O')})/2",
           "E": f"=({t('mw', 'D')}+{t('mw', 'E')})/2", "F": f"=({t('mw', 'E')}+{t('mw', 'F')})/2", "G": f"={t('mwavg', 'E')}*0.75+{t('mwavg', 'F')}*0.25"}
    avg.update({c: f"=({t('mw', PREV[c])}+{t('mw', c)})/2" for c in FC})
    T.row("mwavg", "對外 AI 平均在役 MW（期間）", "MW", avg, "#,##0")
    T.row("len", "期間長度", "年", {"C": 1, "D": 1, "N": 0.75, "O": 0.25, "E": 0.75, "F": 0.25, "G": 1, **{c: 1 for c in FC}}, "0.00")
    add = {"C": f"={t('mw', 'C')}-mw_ye23", "D": f"={t('mw', 'D')}-{t('mw', 'C')}", "N": f"={t('mw', 'N')}-{t('mw', 'C')}", "O": f"={t('mw', 'O')}-{t('mw', 'N')}",
           "E": f"={t('mw', 'E')}-{t('mw', 'D')}", "F": f"={t('mw', 'F')}-{t('mw', 'E')}", "G": f"={t('add', 'E')}+{t('add', 'F')}"}
    add.update({c: f"={t('mw', c)}-{t('mw', PREV[c])}" for c in FC})
    T.row("add", "新增對外 AI MW", "MW", add, "#,##0")
    addT = {c: f"={t('add', c)}/{P('ext_sh')}" for c in ALL}
    addT.update({VC[2023]: f"=mw_ye23/{P('ext_sh')}", VC[2024]: f"={t('addT', 'C')}", VC[2025]: f"={t('addT', 'D')}", VC[2026]: f"={t('addT', 'G')}",
                 **{VC[y]: f"={t('addT', c)}" for y, c in FUT}})
    T.row("addT", "新增 AI 總 MW（含自用＝對外 ÷ 對外比例）", "MW", addT, "#,##0", note="R–AA 欄：依年份（折舊與汰換用）")
    T.sec("2. 新增產能的世代與單位經濟（Tokenomics v21.7 基準成本；自研晶片＝GB300 錨 × IF_NonNV_Hold_Trn，v0.3.1）")
    T.row("vr", "NVIDIA 新增中 VR200 比例", "%", {"F": "=na_vr26", "G": "=na_vr26", "H": "=na_vr27", **{c: "=na_vr28" for c in "IJKLM"}}, "0%")
    mix = lambda pre, fac: {c: f"=cust_sh*{pre}gb300{fac}+(1-cust_sh)*((1-{t('vr', c)})*{pre}gb300+{t('vr', c)}*{pre}vr200)" for c in ["F", "G"] + FC}
    T.row("he_add", "新增 MW 加權 IF_HoldEcon（含自研晶片係數）", "US$m/MW-年", mix("tk_he_", f"*{P('cf_chip')}"), "0.000")
    T.row("ox_add", "新增 MW 加權營運費用（IF_OpexGW）", "US$m/MW-年", mix("tk_ox_", ""), "0.000")
    T.row("cx_add", "新增 MW 加權 IT 資本支出（IF_CapexIT；含自研晶片資本支出係數）", "US$m/MW", mix("tk_cx_", f"*{P('cf_cx')}"), "0.000")
    newc = {"E": f"={t('mw', 'E')}-mw_open", "F": f"={t('mw', 'F')}-mw_open", "G": f"={t('newc', 'F')}", **{c: f"={t('mw', c)}-mw_open" for c in FC}}
    T.row("newc", "期初後新增對外 MW 累計（期末）", "MW", newc, "#,##0")
    for key, src_ in (("he_cum", "he_add"), ("ox_cum", "ox_add")):
        cc = {"E": f"={t('newc', 'E')}*{t(src_, 'F')}", "F": f"={t(key, 'E')}+{t('add', 'F')}*{t(src_, 'F')}", "G": f"={t(key, 'F')}", **{c: f"={t(key, PREV[c])}+{t('add', c)}*{t(src_, c)}" for c in FC}}
        T.row(key, "  Σ 新增 MW × " + ("HoldEcon" if key == "he_cum" else "營運費用"), "US$m", cc, "#,##0")
    T.row("he_new", "新增機隊加權 HoldEcon", "US$m/MW-年", {c: f"=IF({t('newc', c)}>0,{t('he_cum', c)}/{t('newc', c)},{t('he_add', c)})" for c in ["F", "G"] + FC}, "0.000")
    T.row("ox_new", "新增機隊加權營運費用", "US$m/MW-年", {c: f"=IF({t('newc', c)}>0,{t('ox_cum', c)}/{t('newc', c)},{t('ox_add', c)})" for c in ["F", "G"] + FC}, "0.000")
    T.row("rpm_new", "新增產能每 MW 年收入＝新增機隊 HoldEcon × 情境 k", "US$m/MW-年", {c: f"={t('he_new', c)}*{P('kn')}" for c in ["F", "G"] + FC}, "0.000",
          note="6/30 機隊 2,080 MW 三情境同用 rpm_open（k_B）；情境 k 用於 6/30 之後的新產能（含 Q3 新增）", key_fill=True)
    nav = {"F": f"=({t('newc', 'E')}+{t('newc', 'F')})/2", **{c: f"=({t('newc', PREV[c])}+{t('newc', c)})/2" for c in FC}}
    T.row("newavg", "新增產能平均在役 MW（期間）", "MW", nav, "#,##0")
    T.sec("3. AWS AI 損益（收入端不乘 IF_Util：MW-year 計價）")
    ro = P("rpm_open")
    rev = {c: f"={t('mwavg', c)}*{ro}*{t('len', c)}/1000" for c in HIST}
    rev.update({"F": f"=(mw_open*{ro}+{t('newavg', 'F')}*{t('rpm_new', 'F')})*{t('len', 'F')}/1000", "G": f"={t('rev', 'E')}+{t('rev', 'F')}",
                **{c: f"=(mw_open*{ro}+{t('newavg', c)}*{t('rpm_new', c)})/1000" for c in FC}})
    T.row("rev", "AWS AI 營收（命題子列）", "US$bn", rev, bold=True, note="實績欄為分拆估計（期初 k，Derived），不是公司揭露值")
    T.row("rpm_fleet", "全機隊每 MW 年收入", "US$m/MW-年", {c: f"={t('rev', c)}*1000/({t('mwavg', c)}*{t('len', c)})" for c in ALL}, "0.000")
    opx = {c: f"={t('mwavg', c)}/{P('ext_sh')}*(ox_open+rent_mw)*{t('len', c)}/1000" for c in HIST}
    opx.update({c: f"=((mw_open*ox_open+{t('newavg', c)}*{t('ox_new', c)})/{P('ext_sh')}+{t('mwavg', c)}/{P('ext_sh')}*rent_mw)*{t('len', c)}/1000" for c in ["F"] + FC})
    opx["G"] = f"={t('opex', 'E')}+{t('opex', 'F')}"
    T.row("opex", "AI 營運成本（含自用 MW、租用機房租金 C22）", "US$bn", opx)
    T.row("aws_act", "AWS 營收（實績，Q 頁）", "US$bn", {c: f"=«Q.aws.{QM[c]}»" for c in HIST})
    awso = {c: f"={t('aws_act', c)}-{t('rev', c)}" for c in HIST}
    awso.update({"F": f"={t('awso', 'O')}*(1+{t('awso_g', 'F')})", "G": f"={t('awso', 'E')}+{t('awso', 'F')}", **{c: f"={t('awso', PREV[c])}*(1+{t('awso_g', c)})" for c in FC}})
    T.row("awso", "AWS 其他營收（非 AI 雲端）", "US$bn", awso, bold=True, note="實績＝AWS 實績 − AWS AI 分拆估計（殘差，Derived）")
    ag = {"D": f"={t('awso', 'D')}/{t('awso', 'C')}-1", "E": f"={t('awso', 'E')}/{t('awso', 'N')}-1", "F": f"=IF(awso_g0=1,{t('awso_g', 'E')},{t('awso_g', 'D')})", "G": f"={t('awso', 'G')}/{t('awso', 'D')}-1"}
    ag.update(fade(T, "awso_g", P("g_awso_lt"), "F"))
    T.row("awso_g", "  AWS 其他年增", "%", ag, "0.0%", note="起點＝FY2025 年增（D 欄，Derived；In 頁 awso_g0＝0）；9M26 年增（E 欄）受 AI 分拆低估影響偏高，只列不用（awso_g0＝1 時改用，Sens）；收斂到 g_awso_lt")
    T.sec("4. AI 資本支出（現金）→ 折舊（IT 6 年、機房 20 年；次年起提列）")
    T.row("nonai_f", "Q4 26 非 AI 資本支出（零售＋AWS 其他）", "US$bn", {"F": f"=cx_rt*{S.t('rt', 'F')}+cx_awso*{t('awso', 'F')}"})
    T.row("itsh", "AI 資本支出中 IT 比例", "%", {"C": "=tk_cx_h100/(tk_cx_h100+tk_fac*(1-colo_sh))", "D": "=tk_cx_h100/(tk_cx_h100+tk_fac*(1-colo_sh))",
                                               "E": f"={t('itsh', 'F')}", **{c: f"={t('cx_add', c)}/({t('cx_add', c)}+tk_fac*(1-colo_sh))" for c in ["F", "G"] + FC}}, "0.0%")
    cx = {"C": "=«Q.cxn.D»-cx_rt*«Q.rt.D»-cx_awso*" + t("awso", "C"), "D": "=«Q.cxn.E»-cx_rt*«Q.rt.E»-cx_awso*" + t("awso", "D"),
          "E": "=«Q.cxn.O»-cx_rt*«Q.rt.O»-cx_awso*" + t("awso", "E"),
          "F": f"=capex_g26-«Q.cxn.O»-{t('nonai_f', 'F')}", "G": f"={t('cx', 'E')}+{t('cx', 'F')}"}
    cx.update({c: f"={t('addT', c)}*({t('cx_add', c)}+tk_fac*(1-colo_sh))/1000+{t('rfs', c)}" for c in FC})
    cx.update({VC[2023]: f"=«Q.cxn.C»-cx_rt*«Q.rt.C»-cx_awso*(«Q.aws.C»-mw_ye23/2*{P('rpm_open')}/1000)", **{VC[y]: f"={t('cx', c)}" for y, c in VY.items()}})
    T.row("cx", "AI 資本支出（現金，含自用）", "US$bn", cx, bold=True,
          note="實績＝公司資本支出淨額 − 非 AI（零售 cx_rt、AWS 其他 cx_awso）；Q4 26＝指引 220 − 9M 實績 − 非 AI（修正三：FY2026 現金資本支出錨定指引；Q4 新增 MW 部分由已投入的在建資本支出支應）；之後＝新增 MW ×（IT＋自建機房）＋汰換")
    T.row("cx_fml", "C10：Q4 26 AI 資本支出公式值＝新增總 MW ×（IT＋自建機房）", "US$bn", {"F": f"={t('addT', 'F')}*({t('cx_add', 'F')}+tk_fac*(1-colo_sh))/1000"})
    T.row("c10_gap", "C10 對帳（下半年視窗）：（指引 − 上半年實績）÷ 2 −（Q4 公式 AI＋Q4 非 AI）", "US$bn",
          {"F": f"=(capex_g26-«Q.cxn.P»)/2-({t('cx_fml', 'F')}+{t('nonai_f', 'F')})"}, note="情境 B 以 calib_pace_v031.py 解 pace_B 使本列＝0；A／C 情境速度不同，本列不為 0")
    T.row("c10_q4", "對照：Q4 26 指引補足額＝Q4 錨定值 − Q4 公式值（負值＝Q4 公式高於指引隱含）", "US$bn", {"F": f"={t('cx', 'F')}-{t('cx_fml', 'F')}"})
    life = P("tk_life")
    rfs = {c: f"=SUMPRODUCT(--({yrs(T)}={YEAR[c]}-{life}),{T.t('addT', 'R')[:-2]}R»:«{T.id}.addT.AA»)*{t('cx_add', c)}/1000" for c in FC}
    T.row("rfs", "IT 汰換（折舊年限前新增的總 MW × 當年 IT 每 MW）", "US$bn", rfs)
    hcost = "(tk_cx_h100+tk_fac*(1-colo_sh))/1000"
    cxd = {VC[2023]: f"=MIN({t('cx', VC[2023])},{t('addT', VC[2023])}*{hcost})", "C": f"=MIN({t('cx', 'C')},{t('addT', 'C')}*{hcost})",
           "D": f"=MIN({t('cx', 'D')},{t('addT', 'D')}*{hcost})", "E": f"=MIN({t('cx', 'E')},{t('addT', 'E')}*({t('cx_add', 'F')}+tk_fac*(1-colo_sh))/1000)",
           "F": f"={t('cx', 'F')}", "G": f"={t('cx_dep', 'E')}+{t('cx_dep', 'F')}", **{c: f"={t('cx', c)}" for c in FC}}
    cxd.update({VC[y]: f"={t('cx_dep', c)}" for y, c in VY.items()})
    T.row("cx_dep", "AI 資本支出：折舊基礎（實績期以新增在役 MW × 每 MW 成本為上限）", "US$bn", cxd,
          note="實績期以新增在役 MW × 每 MW 成本為上限；超出部分見下列在建（修正三：延後轉入在役後提列折舊）")
    cipd = {c: f"={t('cx', c)}-{t('cx_dep', c)}" for c in ["C", "D", "E", "F", "G"] + FC}
    cipd.update({VC[y]: f"={t('cx', VC[y])}-{t('cx_dep', VC[y])}" for y in range(2023, 2033)})
    T.row("cx_cip", "  在建（CIP）：實績期超出新增在役 MW 成本的 AI 資本支出", "US$bn", cipd,
          note="修正三：視為在建，延後 cip_lag 年轉入在役後提列折舊（仍計入 AI 投入資本）；AWS 分部折舊 2024→2025→1H26 年化 13.3→21.5→30.7 支持此處理")
    cipit = {VC[y]: f"={t('cx_cip', VC[y])}*{t('itsh', 'C' if y <= 2025 else 'E')}" for y in range(2023, 2033)}
    T.row("cip_it", "  在建中 IT 設備（依支出年份）", "US$bn", cipit)
    def shift(key):
        return {VC[y]: f"=SUMPRODUCT(--({yrs(T)}={y}-{P('cip_lag')}),«{T.id}.{key}.R»:«{T.id}.{key}.AA»)" for y in range(2023, 2033)}
    T.row("cip_it_s", "  在建轉入在役：IT（依轉入年份）", "US$bn", shift("cip_it"))
    cipfs = {VC[y]: f"=SUMPRODUCT(--({yrs(T)}={y}-{P('cip_lag')}),«{T.id}.cx_cip.R»:«{T.id}.cx_cip.AA»)-{t('cip_it_s', VC[y])}" for y in range(2023, 2033)}
    T.row("cip_fac_s", "  在建轉入在役：機房（依轉入年份）", "US$bn", cipfs)
    cxit = {"C": f"={t('cx_dep', 'C')}*{t('itsh', 'C')}", "D": f"={t('cx_dep', 'D')}*{t('itsh', 'D')}", "E": f"={t('cx_dep', 'E')}*{t('itsh', 'E')}", "F": f"={t('cx_dep', 'F')}*{t('itsh', 'F')}",
            "G": f"={t('cxit', 'E')}+{t('cxit', 'F')}", **{c: f"=({t('cx_dep', c)}-{t('rfs', c)})*{t('itsh', c)}+{t('rfs', c)}" for c in FC}}
    cxit.update({VC[2023]: f"={t('cx_dep', VC[2023])}*{t('itsh', 'C')}", **{VC[y]: f"={t('cxit', c)}" for y, c in VY.items()}})
    T.row("cxit", "  其中 IT 設備", "US$bn", cxit)
    T.row("cxfac", "  其中機房（自建部分）", "US$bn", {**{VC[y]: f"={t('cx_dep', VC[y])}-{t('cxit', VC[y])}" for y in range(2023, 2033)},
                                               **{c: f"={t('cx_dep', c)}-{t('cxit', c)}" for c in ["C", "D", "E", "F", "G"] + FC}})
    T.row("cx_sh", "AI 資本支出占公司資本支出", "%", {c: f"={t('cx', c)}/«{eid}C.cx.{c}»" for c in ["C", "D", "E", "F", "G"] + FC}, "0.0%")
    rng = lambda key: f"«{T.id}.{key}.R»:«{T.id}.{key}.AA»"
    hrow = REG[(T.id, "hdr")][1]
    def daspec(key, f):
        d = {VC[y]: f.format(y=f"{T.ws.title}!{VC[y]}${hrow}") for y in range(2023, 2033)}
        d.update({c: f"={t(key, VC[y])}" for y, c in VY.items()})
        d.update({"N": f"={t(key, 'D')}*0.75", "O": f"={t(key, 'D')}*0.25", "E": f"={t(key, 'G')}*0.75", "F": f"={t(key, 'G')}*0.25"})
        return d
    T.row("da_it", "AI IT 折舊（含在建轉入）", "US$bn", daspec("da_it", f"=(SUMPRODUCT(--({yrs(T)}<{{y}}),--({yrs(T)}>={{y}}-{life}),{rng('cxit')})+SUMPRODUCT(--({yrs(T)}<{{y}}),--({yrs(T)}>={{y}}-{life}),{rng('cip_it_s')}))/{life}"))
    T.row("da_fac", "AI 機房折舊（含在建轉入）", "US$bn", daspec("da_fac", f"=(SUMPRODUCT(--({yrs(T)}<{{y}}),{rng('cxfac')})+SUMPRODUCT(--({yrs(T)}<{{y}}),{rng('cip_fac_s')}))/{P('fac_life')}"))
    dd = {c: f"={t('da_it', c)}+{t('da_fac', c)}" for c in ALL + [VC[y] for y in range(2023, 2033)]}
    T.row("da", "AI 折舊合計（對外 → AWS；自用 → 零售分部）", "US$bn", dd, bold=True)
    ppe = {VC[y]: f"=SUM(«{T.id}.cx.R»:{t('cx', VC[y])})-SUM(«{T.id}.da.R»:{t('da', VC[y])})" for y in range(2023, 2033)}
    ppe.update({c: f"={t('ppe', VC[y])}" for y, c in VY.items()})
    T.row("ppe", "AI 淨 PP&E（期末）", "US$bn", ppe)
    T.sec("5. AWS AI 營業利益、自用 AI 成本、AWS 其他")
    T.row("ebitda", "AWS AI EBITDA＝營收 − 對外比例 × 營運成本", "US$bn", {c: f"={t('rev', c)}-{P('ext_sh')}*{t('opex', c)}" for c in ALL})
    T.row("ebit", "AWS AI 營業利益＝EBITDA − 對外比例 × AI 折舊", "US$bn", {c: f"={t('ebitda', c)}-{P('ext_sh')}*{t('da', c)}" for c in ALL}, bold=True)
    T.row("ebit_m", "  AWS AI 營業利益率", "%", {c: f"={t('ebit', c)}/{t('rev', c)}" for c in ALL}, "0.0%")
    rd = f"IF({P('internal_ai_reading')}=2,1,0)"
    shd = {c: 0 for c in HIST}
    shd.update({c: f"={rd}*(1-{P('ext_sh')})/{P('ext_sh')}*(mw_open*he_open+{t('newavg', c)}*{t('he_new', c)})*{t('len', c)}/1000" for c in ["F"] + FC})
    shd["G"] = f"={t('shadow', 'E')}+{t('shadow', 'F')}"
    T.row("shadow", "自用 AI 影子效益（讀法 2：自用 MW × HoldEcon × k＝1；讀法 1 為 0；非財報營收）", "US$bn", shd,
          note="In 頁 internal_ai_reading＝2（主值）；只進評價口徑（NOPAT、EPS 評價口徑、分部加總零售本業）；實績欄為 0")
    T.row("int_cost", "自用 AI 成本（營運＋折舊；分攤到北美／國際，財報口徑）", "US$bn", {c: f"=(1-{P('ext_sh')})*({t('opex', c)}+{t('da', c)})" for c in ALL})
    T.row("aws_oi_act", "AWS 營業利益（實績，Q 頁）", "US$bn", {c: f"=«Q.oi_aws.{QM[c]}»" for c in HIST})
    oe = {c: f"={t('aws_oi_act', c)}-{t('ebit', c)}" for c in HIST}
    oe.update({"F": f"={t('awso', 'F')}*{t('awso_m', 'F')}", "G": f"={t('awso_ebit', 'E')}+{t('awso_ebit', 'F')}", **{c: f"={t('awso', c)}*{t('awso_m', c)}" for c in FC}})
    T.row("awso_ebit", "AWS 其他營業利益", "US$bn", oe, note="實績＝AWS 營業利益 − AWS AI 營業利益（殘差）")
    om = {c: f"={t('awso_ebit', c)}/{t('awso', c)}" for c in HIST}
    om.update({"F": f"=({t('awso_ebit', 'O')}+{t('awso_ebit', 'E')})/({t('awso', 'O')}+{t('awso', 'E')})", "G": f"={t('awso_ebit', 'G')}/{t('awso', 'G')}", **{c: f"={t('awso_m', 'F')}" for c in FC}})
    T.row("awso_m", "  AWS 其他營業利益率", "%", om, "0.0%", note="Q4 26 起固定為近四季（2025Q4＋9M26）殘差利益率（Derived）")
    T.row("da_awso_r", "AWS 非 AI 折舊 ÷ AWS 其他營收（FY2025 校準）", "%", {"D": f"=(«Q.da_aws.E»-{P('ext_sh')}*{t('da', 'D')})/{t('awso', 'D')}"}, "0.0%")
    daws = {c: f"=«Q.da_aws.{QM[c]}»" for c in HIST}
    daws.update({c: f"={P('ext_sh')}*{t('da', c)}+{t('da_awso_r', 'D')}*{t('awso', c)}" for c in ["F"] + FC})
    daws["G"] = f"={t('da_aws', 'E')}+{t('da_aws', 'F')}"
    T.row("da_aws", "AWS 折舊攤銷（分部）", "US$bn", daws)
    T.sec("6. 命題指標")
    roic = {"G": f"={t('ebit', 'G')}*(1-tax)/({P('ext_sh')}*({t('ppe', 'D')}+{t('ppe', 'G')})/2)",
            **{c: f"={t('ebit', c)}*(1-tax)/({P('ext_sh')}*({t('ppe', PREV[c])}+{t('ppe', c)})/2)" for c in FC}}
    T.row("roic", "對外 AI ROIC（稅後營業利益 ÷ 對外分攤的平均 AI 淨 PP&E）", "%", roic, "0.0%", key_fill=True, note="C11：分子分母同口徑（對外）")
    keff = {c: f"={t('rev', c)}*1000/(mw_open*he_open+{t('newavg', c)}*{t('he_new', c)})" for c in FC}
    T.row("k_eff", "機隊平均 k（營收 ÷ 持有成本）", "x", keff, "0.000")
    kbe = {c: f"={t('k_eff', c)}*({t('rev', c)}+{P('wacc')}*{P('ext_sh')}*({t('ppe', PREV[c])}+{t('ppe', c)})/2/(1-tax)-{t('ebit', c)})/{t('rev', c)}" for c in FC}
    T.row("k_be", "打平 k（對外 AI ROIC＝WACC 所需的機隊平均 k）", "x", kbe, "0.000")
    T.row("pq_chk", "上限檢查：全機隊每 MW 年收入 ÷ Tokenomics P×Q（VR200）", "%", {c: f"={t('rpm_fleet', c)}/tk_pq" for c in ALL}, "0.0%", note="上限，不得作預測；只作檢查")
    T.close()

    # ---------------- Seg（二）：AWS 與合併 ----------------
    S.sec("2. AWS（命題子列＋AWS 其他）與合併營收")
    S.row("aws_ai", "AWS AI（命題子列）", "US$bn", {c: f"={T.t('rev', c)}" for c in ALL}, note="Thesis 頁")
    S.row("aws_o", "AWS 其他", "US$bn", {c: f"={T.t('awso', c)}" for c in ALL}, note="Thesis 頁")
    aw = {c: f"=«Q.aws.{QM[c]}»" for c in HIST}
    aw.update({c: f"={S.t('aws_ai', c)}+{S.t('aws_o', c)}" for c in ["F", "G"] + FC})
    S.row("aws", "AWS", "US$bn", aw, bold=True, note="官方分部＝AWS 營收線")
    tot = {c: f"=«Q.tot.{QM[c]}»" for c in HIST}
    tot.update({c: f"={S.t('rt', c)}+{S.t('aws', c)}" for c in ["F", "G"] + FC})
    S.row("tot", "合併營收（實績＝公司揭露；預測＝營收線加總）", "US$bn", tot, bold=True)
    S.row("chk_rt", f"檢查（情境 {s}）：零售營收線合計 − (North America＋International)", "US$bn", {c: f"={S.t('rt', c)}-({S.t('na', c)}+{S.t('intl', c)})" for c in ALL}, check=True)
    S.row("chk_aws", f"檢查（情境 {s}）：AWS AI＋AWS 其他 − AWS", "US$bn", {c: f"={S.t('aws_ai', c)}+{S.t('aws_o', c)}-{S.t('aws', c)}" for c in ALL}, check=True)
    S.row("chk_seg", f"檢查（情境 {s}）：分部合計 − 合併營收", "US$bn", {c: f"={S.t('na', c)}+{S.t('intl', c)}+{S.t('aws', c)}-{S.t('tot', c)}" for c in ALL}, check=True)
    S.row("g_tot", "合併營收年增", "%", {c: f"={S.t('tot', c)}/{S.t('tot', PREV.get(c, 'D'))}-1" for c in ["G"] + FC}, "0.0%")
    S.row("aws_sh", "AWS 占合併營收", "%", {c: f"={S.t('aws', c)}/{S.t('tot', c)}" for c in ALL}, "0.0%")
    S.row("ai_sh", "AWS AI 占 AWS", "%", {c: f"={S.t('aws_ai', c)}/{S.t('aws', c)}" for c in ALL}, "0.0%")
    S.close()

    # ---------------- PL ----------------
    L = Blk(f"{eid}P", sheets["PL"], f"PL｜{nm}：分部營收 → 分部營業利益 → 合併營業利益 → 業外、稅 → 淨利 → EPS")
    l = L.t
    L.sec("1. 營收（10-K 營收線與官方分部）")
    for k, lab, _ in LINES:
        L.row("rev_" + k, lab, "US$bn", {c: f"={S.t('rev_' + k, c)}" for c in ALL})
    L.row("awsl", "AWS（營收線＝分部）", "US$bn", {c: f"={S.t('aws', c)}" for c in ALL})
    L.row("na", "North America", "US$bn", {c: f"={S.t('na', c)}" for c in ALL}, bold=True)
    L.row("intl", "International", "US$bn", {c: f"={S.t('intl', c)}" for c in ALL}, bold=True)
    L.row("aws", "AWS", "US$bn", {c: f"={S.t('aws', c)}" for c in ALL}, bold=True)
    L.row("aws_ai", "AWS AI（命題子列）", "US$bn", {c: f"={S.t('aws_ai', c)}" for c in ALL})
    L.row("aws_o", "AWS 其他", "US$bn", {c: f"={S.t('aws_o', c)}" for c in ALL})
    L.row("tot", "合併營收", "US$bn", {c: f"={S.t('tot', c)}" for c in ALL}, bold=True)
    L.row("chk_rt", f"檢查（情境 {s}）：零售營收線合計 − (North America＋International)", "US$bn",
          {c: "=" + "+".join(l("rev_" + k, c) for k, _, _ in LINES) + f"-({l('na', c)}+{l('intl', c)})" for c in ALL}, check=True)
    L.row("chk_aws", f"檢查（情境 {s}）：AWS AI＋AWS 其他 − AWS", "US$bn", {c: f"={l('aws_ai', c)}+{l('aws_o', c)}-{l('aws', c)}" for c in ALL}, check=True)
    L.row("chk_ln", f"檢查（情境 {s}）：營收線合計（含 AWS）− 合併營收", "US$bn",
          {c: "=" + "+".join(l("rev_" + k, c) for k, _, _ in LINES) + f"+{l('awsl', c)}-{l('tot', c)}" for c in ALL}, check=True)
    L.row("chk_seg", f"檢查（情境 {s}）：分部合計 − 合併營收", "US$bn", {c: f"={l('na', c)}+{l('intl', c)}+{l('aws', c)}-{l('tot', c)}" for c in ALL}, check=True)
    L.sec("2. 分部營業利益（廣告利潤率 m_ad 計入北美／國際；零售本業利益率＝殘差；自用 AI 成本按北美比例分攤）")
    for seg, nm_, step, sh_ in (("na", "North America", "step_na", f"{S.t('na_sh', '{c}')}"), ("in", "International", "step_in", f"(1-{S.t('na_sh', '{c}')})")):
        adk = "ad_na" if seg == "na" else "ad_in"
        segk = "na" if seg == "na" else "intl"
        L.row("ads_" + seg, f"{nm_}：廣告利潤（廣告營收 × m_ad）", "US$bn", {c: f"={S.t(adk, c)}*{P('m_ad')}" for c in ALL})
        L.row("core_" + seg, f"{nm_}：零售本業營收（分部 − 廣告）", "US$bn", {c: f"={S.t(segk, c)}-{S.t(adk, c)}" for c in ALL})
        L.row("int_" + seg, f"{nm_}：自用 AI 成本淨額分攤", "US$bn", {c: f"={T.t('int_cost', c)}*" + sh_.replace("{c}", c) for c in ALL})
        oi = {c: f"=«Q.oi_{seg}.{QM[c]}»" for c in HIST}
        oi.update({c: f"={l('ads_' + seg, c)}+{l('core_' + seg, c)}*{l('cm_' + seg, c)}-{l('int_' + seg, c)}" for c in ["F"] + FC})
        oi["G"] = f"={l('oi_' + seg, 'E')}+{l('oi_' + seg, 'F')}"
        cm = {c: f"=({l('oi_' + seg, c)}-{l('ads_' + seg, c)}+{l('int_' + seg, c)})/{l('core_' + seg, c)}" for c in HIST + ["G"]}
        cm["F"] = f"={l('cm_' + seg, 'O')}+({l('cm_' + seg, 'E')}-{l('cm_' + seg, 'N')})"
        cm.update({c: f"={l('cm_' + seg, PREV[c])}+{P(step)}" for c in "HIJK"})
        cm.update({c: f"={l('cm_' + seg, PREV[c])}" for c in "LM"})
        L.row("cm_" + seg, f"{nm_}：零售本業利益率（實績＝殘差，Derived）", "%", cm, "0.00%", note="Q4 26＝Q4 25＋(9M26 − 9M25)；2027–2030 每年＋In 頁 step；之後持平")
        L.row("oi_" + seg, f"{nm_} 營業利益", "US$bn", oi, bold=True)
    L.row("oi_aws_ai", "AWS AI 營業利益（命題子列）", "US$bn", {c: f"={T.t('ebit', c)}" for c in ALL})
    L.row("oi_aws_o", "AWS 其他營業利益", "US$bn", {c: f"={T.t('awso_ebit', c)}" for c in ALL})
    oa = {c: f"=«Q.oi_aws.{QM[c]}»" for c in HIST}
    oa.update({c: f"={l('oi_aws_ai', c)}+{l('oi_aws_o', c)}" for c in ["F", "G"] + FC})
    L.row("oi_aws", "AWS 營業利益", "US$bn", oa, bold=True)
    un = {c: f"=«Q.unal.{QM[c]}»" for c in HIST}
    un.update({c: 0 for c in ["F"] + FC}); un["G"] = f"={l('unal', 'E')}+{l('unal', 'F')}"
    L.row("unal", "未分攤（公司層；Amazon 為 0）", "US$bn", un)
    oit = {c: f"=«Q.oi.{QM[c]}»" for c in HIST}
    oit.update({c: f"={l('oi_na', c)}+{l('oi_in', c)}+{l('oi_aws', c)}+{l('unal', c)}" for c in ["F", "G"] + FC})
    L.row("oi", "合併營業利益", "US$bn", oit, bold=True)
    L.row("chk_oi", f"檢查（情境 {s}）：分部營業利益合計 − 合併營業利益", "US$bn", {c: f"={l('oi_na', c)}+{l('oi_in', c)}+{l('oi_aws', c)}+{l('unal', c)}-{l('oi', c)}" for c in ALL}, check=True)
    L.row("shadow", "自用 AI 影子效益（讀法 2；非財報營收；按北美比例分攤到北美／國際）", "US$bn", {c: f"={T.t('shadow', c)}" for c in ALL}, note="營收列與分部對帳不動")
    L.row("oi_v", "營業利益（評價口徑＝合併營業利益＋自用 AI 影子效益）", "US$bn", {c: f"={l('oi', c)}+{l('shadow', c)}" for c in ALL})
    for k, a, b_ in (("m_na", "oi_na", "na"), ("m_in", "oi_in", "intl"), ("m_aws", "oi_aws", "aws"), ("m_ai", "oi_aws_ai", "aws_ai"), ("m_oi", "oi", "tot")):
        lab = {"m_na": "North America 營業利益率", "m_in": "International 營業利益率", "m_aws": "AWS 營業利益率", "m_ai": "AWS AI 營業利益率", "m_oi": "合併營業利益率"}[k]
        L.row(k, lab, "%", {c: f"={l(a, c)}/{l(b_, c)}" for c in ALL}, "0.0%", key_fill=(k == "m_aws"))
    L.sec("3. 分部折舊攤銷（PP&E）與 EBITDA（分部加總評價用）")
    L.row("da_rt_r", "零售非 AI 折舊 ÷ 零售營收（FY2025 校準）", "%", {"D": f"=(«Q.da_na.E»+«Q.da_in.E»-(1-{P('ext_sh')})*{T.t('da', 'D')})/{S.t('rt', 'D')}"}, "0.00%")
    for seg, sh_ in (("na", f"{S.t('na_sh', '{c}')}"), ("in", f"(1-{S.t('na_sh', '{c}')})")):
        d = {c: f"=«Q.da_{seg}.{QM[c]}»" for c in HIST}
        d.update({c: f"=({l('da_rt_r', 'D')}*{S.t('rt', c)}+(1-{P('ext_sh')})*{T.t('da', c)})*" + sh_.replace("{c}", c) for c in ["F"] + FC})
        d["G"] = f"={l('da_' + seg, 'E')}+{l('da_' + seg, 'F')}"
        L.row("da_" + seg, ("North America" if seg == "na" else "International") + " 折舊攤銷", "US$bn", d)
    L.row("da_aws", "AWS 折舊攤銷", "US$bn", {c: f"={T.t('da_aws', c)}" for c in ALL})
    L.row("da", "分部折舊攤銷合計（PP&E；不含影音內容與營業租賃資產）", "US$bn", {c: f"={l('da_na', c)}+{l('da_in', c)}+{l('da_aws', c)}" for c in ALL})
    L.row("e_ads", "EBITDA：廣告（＝廣告利潤，未分攤折舊）", "US$bn", {c: f"={l('ads_na', c)}+{l('ads_in', c)}" for c in ALL})
    L.row("e_rt", "EBITDA：零售本業（北美＋國際 營業利益＋折舊 − 廣告利潤；含自用 AI 成本）", "US$bn", {c: f"={l('oi_na', c)}+{l('oi_in', c)}+{l('da_na', c)}+{l('da_in', c)}-{l('e_ads', c)}" for c in ALL})
    L.row("e_awso", "EBITDA：AWS 其他", "US$bn", {c: f"={l('oi_aws_o', c)}+{l('da_aws', c)}-{P('ext_sh')}*{T.t('da', c)}" for c in ALL})
    L.row("e_ai", "EBITDA：AWS AI", "US$bn", {c: f"={T.t('ebitda', c)}" for c in ALL})
    L.row("ebitda", "EBITDA（營業利益＋分部折舊攤銷）", "US$bn", {c: f"={l('oi', c)}+{l('da', c)}" for c in ALL})
    L.sec("4. 業外、稅、淨利、EPS")
    ix = {c: f"=«Q.intx.{QM[c]}»" for c in HIST}
    ix.update({c: f"=«{eid}C.intx.{c}»" for c in ["F"] + FC}); ix["G"] = f"={l('intx', 'E')}+{l('intx', 'F')}"
    L.row("intx", "利息費用", "US$bn", ix)
    ii = {c: f"=«{eid}C.inti.{c}»" for c in ["F"] + FC}
    L.row("inti", "利息收入（預測；實績含在業外其他）", "US$bn", ii)
    no = {c: f"=«Q.nonop.{QM[c]}»" for c in ["C", "D", "N", "O", "E"]}
    no.update({c: 0 for c in ["F"] + FC}); no["G"] = f"={l('nonop', 'E')}+{l('nonop', 'F')}"
    L.row("nonop", "業外其他（含利息收入與 Anthropic 等持股評價損益；預測期設 0）", "US$bn", no, note="持股價值另在 Val 頁加入（不重複計算）")
    pt = {c: f"=«Q.pretax.{QM[c]}»" for c in HIST}
    pt.update({c: f"={l('oi', c)}-{l('intx', c)}+{l('inti', c)}+{l('nonop', c)}" for c in ["F"] + FC}); pt["G"] = f"={l('pretax', 'E')}+{l('pretax', 'F')}"
    L.row("pretax", "稅前淨利", "US$bn", pt)
    tx = {c: f"=«Q.taxx.{QM[c]}»" for c in HIST}
    tx.update({c: f"={l('pretax', c)}*tax" for c in ["F"] + FC}); tx["G"] = f"={l('taxx', 'E')}+{l('taxx', 'F')}"
    L.row("taxx", "所得稅", "US$bn", tx)
    ni = {c: f"=«Q.ni.{QM[c]}»" for c in HIST}
    ni.update({c: f"={l('pretax', c)}-{l('taxx', c)}" for c in ["F"] + FC}); ni["G"] = f"={l('ni', 'E')}+{l('ni', 'F')}"
    L.row("ni", "淨利", "US$bn", ni, bold=True)
    shs = {"C": "=«Q.shd.D»", "D": "=«Q.shd.E»", "E": "=«Q.shd.O»", "F": f"=«Q.shd.M»*(1+{P('sh_g')}*0.25)", "G": f"=(3*{l('sh', 'E')}+{l('sh', 'F')})/4"}
    shs.update({c: f"={l('sh', PREV[c])}*(1+{P('sh_g')})" for c in FC})
    L.row("sh", "稀釋加權股數", "十億股", shs, "#,##0.000", note="無回購；SBC 稀釋每年 sh_g")
    L.row("ni_v", "淨利（評價口徑＝淨利＋影子效益 ×（1 − 稅率））", "US$bn", {c: f"={l('ni', c)}+{l('shadow', c)}*(1-tax)" for c in ALL})
    L.row("eps", "EPS（評價口徑，稀釋）", "US$", {c: f"={l('ni_v', c)}/{l('sh', c)}" for c in ["C", "D", "E", "G"] + FC}, "0.00", bold=True, key_fill=True)
    L.row("eps_fin", "EPS（財報口徑，稀釋）", "US$", {c: f"={l('ni', c)}/{l('sh', c)}" for c in ["C", "D", "E", "G"] + FC}, "0.00")
    L.row("eps_x", "EPS（不含業外其他，近似：扣除業外其他 ×（1 − 稅率））", "US$", {c: f"=({l('ni', c)}-{l('nonop', c)}*(1-tax))/{l('sh', c)}" for c in ["C", "D", "E", "G"] + FC}, "0.00",
          note="對照 S&P 調整後 EPS；9M26 業外含 Anthropic 評價利益約 68")
    L.row("eps_g", "EPS 年增", "%", {c: f"={l('eps', c)}/{l('eps', PREV[c])}-1" for c in FC}, "0.0%")
    L.close()

    # ---------------- CF ----------------
    C = Blk(f"{eid}C", sheets["CF"], f"CF｜{nm}：營運現金 − 資本支出 − 租賃本金 → 自由現金流；股利、回購、借款與淨現金")
    k_ = C.t
    CF5 = ["C", "D", "E", "F", "G"] + FC
    C.row("ni", "淨利", "US$bn", {c: f"={l('ni', c)}" for c in CF5})
    C.row("da", "分部折舊攤銷（PP&E）", "US$bn", {c: f"={l('da', c)}" for c in CF5})
    sb = {c: f"=«Q.sbc.{QM[c]}»" for c in ["C", "D", "E"]}
    sb.update({c: f"=sbc_r*{S.t('tot', c)}" for c in ["F"] + FC}); sb["G"] = f"={k_('sbc', 'E')}+{k_('sbc', 'F')}"
    C.row("sbc", "股份基礎給付（加回）", "US$bn", sb)
    ot = {c: f"=«Q.cfo.{QM[c]}»-{k_('ni', c)}-{k_('da', c)}-{k_('sbc', c)}" for c in ["C", "D", "E"]}
    ot.update({c: f"=oth_r*{S.t('tot', c)}" for c in FC}); ot["F"] = f"=oth_q4*{S.t('tot', 'F')}"; ot["G"] = f"={k_('oth', 'E')}+{k_('oth', 'F')}"
    C.row("oth", "其他營運項目（營運資金、影音內容與租賃現金支出、非現金業外）", "US$bn", ot, note="Q4 26＝oth_q4（2025Q4 同期比例，C14）× 營收；FY2027 起 oth_r（已含營運資金）；實績含持股評價利益的扣除")
    cf = {c: f"=«Q.cfo.{QM[c]}»" for c in ["C", "D", "E"]}
    cf.update({c: f"={k_('ni', c)}+{k_('da', c)}+{k_('sbc', c)}+{k_('oth', c)}" for c in ["F"] + FC}); cf["G"] = f"={k_('cfo', 'E')}+{k_('cfo', 'F')}"
    C.row("cfo", "營運現金流", "US$bn", cf, bold=True)
    C.row("cx_ai", "資本支出：AI（Thesis）", "US$bn", {c: f"={T.t('cx', c)}" for c in CF5})
    co = {c: f"=cx_awso*{S.t('aws_o', c)}" for c in CF5}; co["G"] = f"={k_('cx_awso', 'E')}+{k_('cx_awso', 'F')}"
    C.row("cx_awso", "資本支出：AWS 其他（cx_awso × AWS 其他營收）", "US$bn", co)
    cr = {c: f"=«Q.cxn.{QM[c]}»-{k_('cx_ai', c)}-{k_('cx_awso', c)}" for c in ["C", "D", "E"]}
    cr.update({c: f"=cx_rt*{S.t('rt', c)}" for c in ["F"] + FC}); cr["G"] = f"={k_('cx_rt', 'E')}+{k_('cx_rt', 'F')}"
    C.row("cx_rt", "資本支出：零售（cx_rt × 零售營收）", "US$bn", cr)
    cxx = {c: f"=«Q.cxn.{QM[c]}»" for c in ["C", "D", "E"]}
    cxx.update({c: f"={k_('cx_ai', c)}+{k_('cx_awso', c)}+{k_('cx_rt', c)}" for c in ["F"] + FC}); cxx["G"] = f"={k_('cx', 'E')}+{k_('cx', 'F')}"
    C.row("cx", "資本支出（現金，淨額）", "US$bn", cxx, bold=True)
    C.row("lease", "融資租賃本金償還", "US$bn", {"C": "=«Q.lpr.D»", "D": "=«Q.lpr.E»", "E": "=«Q.lpr.O»", "F": "=lease_pr/4", "G": f"={k_('lease', 'E')}+{k_('lease', 'F')}", **{c: "=lease_pr" for c in FC}},
          note="9M26 實績 1.296（新聞稿現金流量表）；Q4 26＝全年 ÷ 4")
    C.row("fcf", "自由現金流（營運現金流 − 資本支出 − 融資租賃本金）", "US$bn", {c: f"={k_('cfo', c)}-{k_('cx', c)}-{k_('lease', c)}" for c in CF5}, bold=True, key_fill=True)
    C.row("div", "股利", "US$bn", {c: f"=dps*{l('sh', c)}" for c in ["C", "D", "E", "G"] + FC} | {"F": "=0"}, note="Amazon 未配息")
    C.row("bb", "回購", "US$bn", {c: "=0" for c in CF5}, note="2023–9M26 無回購；模型不設回購（無可減的回購）")
    C.row("fcfd", "自由現金流扣股利", "US$bn", {c: f"={k_('fcf', c)}-{k_('div', c)}" for c in CF5}, bold=True)
    C.row("inv", "策略投資（OpenAI、Anthropic 等）", "US$bn", {"E": "=«Q.acq.O»", "F": "=inv_q4", "G": f"={k_('inv', 'E')}+{k_('inv', 'F')}", **{c: 0 for c in FC}})
    C.row("cash_o", "期初現金（Q4 26 起；2026-09-30 實績）", "US$bn", {"F": "=cash_q3", "G": "=cash_q3", **{c: f"={k_('cash', PREV[c])}" for c in FC}})
    C.row("pre", "借款前現金", "US$bn", {c: f"={k_('cash_o', c)}+{k_('fcfd', c)}-{k_('bb', c)}-{k_('inv', c)}" for c in ["F"] + FC})
    bo = {c: f"=MAX(0,min_cash-{k_('pre', c)})" for c in ["F"] + FC}; bo["G"] = f"={k_('borrow', 'F')}"
    C.row("borrow", "新增借款（現金低於最低現金時）", "US$bn", bo, note="到期債務假設全數再融資（Assumed）")
    ca = {"D": "=«Q.cash.E»", "E": "=«Q.cash.M»", "F": f"={k_('pre', 'F')}+{k_('borrow', 'F')}", "G": f"={k_('cash', 'F')}", **{c: f"={k_('pre', c)}+{k_('borrow', c)}" for c in FC}}
    C.row("cash", "期末現金與有價證券", "US$bn", ca)
    de = {"D": "=«Q.debt.E»", "E": "=«Q.debt.M»", "F": f"=debt_q3+{k_('borrow', 'F')}", "G": f"={k_('debt', 'F')}", **{c: f"={k_('debt', PREV[c])}+{k_('borrow', c)}" for c in FC}}
    C.row("debt", "期末債務面額", "US$bn", de)
    C.row("nd", "淨負債（債務 − 現金；負值＝淨現金）", "US$bn", {c: f"={k_('debt', c)}-{k_('cash', c)}" for c in ["D", "E", "F", "G"] + FC})
    C.row("intx", "利息費用（既有債務 × kd_book＋新增債務 × kd_new）", "US$bn",
          {"F": "=debt_q3*kd_book*0.25", **{c: f"=debt_q3*kd_book+({k_('debt', PREV[c])}-debt_q3)*kd_new" for c in FC}})
    C.row("inti", "利息收入（期初現金 × y_cash）", "US$bn", {"F": "=cash_q3*y_cash*0.25", **{c: f"={k_('cash', PREV[c])}*y_cash" for c in FC}})
    C.row("gate", "收支關卡（本情境）：自由現金流扣股利 FY2026–FY2030 合計", "US$bn", {"G": f"=SUM({k_('fcfd', 'G')}:{k_('fcfd', 'K')})"}, bold=True, key_fill=True)
    C.row("borrow5", "FY2026–FY2030 新增借款合計（不含已發行債券）", "US$bn", {"G": f"=SUM({k_('borrow', 'G')}:{k_('borrow', 'K')})"})
    C.row("chk_gate", f"檢查（情境 {s}）：收支關卡＝首年起五年自由現金流扣股利（逐欄加總）", "US$bn",
          {"G": f"={k_('gate', 'G')}-(" + "+".join(k_('fcfd', c) for c in "GHIJK") + ")"}, check=True)
    C.close()

    # ---------------- Val ----------------
    V = Blk(f"{eid}V", sheets["Val"], f"Val｜{nm}：DCF（FCFF，WACC）與分部加總（SOTP，同業 EV/EBITDA）")
    v = V.t
    V.sec("DCF（基準日 2026-12-31；FY2027–FY2032 期中折現；永續價值＝FY2033 NOPAT ×（1 − g ÷ RONIC）÷（WACC − g））")
    V.row("nopat", "NOPAT＝營業利益（評價口徑，含自用 AI 影子效益）×（1 − 稅率）", "US$bn", {c: f"={l('oi_v', c)}*(1-tax)" for c in FC})
    V.row("da", "＋分部折舊攤銷", "US$bn", {c: f"={l('da', c)}" for c in FC})
    V.row("cx", "− 資本支出", "US$bn", {c: f"={k_('cx', c)}" for c in FC})
    V.row("lease", "− 新增融資租賃資產（以融資租賃本金代理；非現金取得，修正三）", "US$bn", {c: f"={k_('lease', c)}" for c in FC})
    V.row("oth", "＋其他營運項目", "US$bn", {c: f"={k_('oth', c)}" for c in FC})
    V.row("ufcf", "FCFF（SBC 不加回，視同現金成本）", "US$bn", {c: f"={v('nopat', c)}+{v('da', c)}-{v('cx', c)}-{v('lease', c)}+{v('oth', c)}" for c in FC}, bold=True)
    V.row("df", "折現因子（期中）", "x", {c: f"=1/(1+{P('wacc')})^({YEAR[c]}-2026-0.5)" for c in FC}, "0.0000")
    V.row("pv", "現值", "US$bn", {c: f"={v('ufcf', c)}*{v('df', c)}" for c in FC})
    V.row("spv", "明確期現值合計", "US$bn", {"C": f"=SUM({v('pv', 'H')}:{v('pv', 'M')})"})
    V.row("tv", "永續價值（2032 年底）", "US$bn", {"C": f"={v('nopat', 'M')}*(1+{P('tg')})*(1-{P('tg')}/{P('ronic')})/({P('wacc')}-{P('tg')})"})
    V.row("pvtv", "永續價值現值", "US$bn", {"C": f"={v('tv', 'C')}/(1+{P('wacc')})^6"})
    V.row("ev", "企業價值（EV）", "US$bn", {"C": f"={v('spv', 'C')}+{v('pvtv', 'C')}"})
    V.row("tvsh", "永續價值占 EV", "%", {"C": f"={v('pvtv', 'C')}/{v('ev', 'C')}"}, "0.0%")
    V.row("nd26", "− 淨負債（2026 年底，含融資租賃與融資義務）", "US$bn", {"C": f"={k_('nd', 'G')}+fin_debt"})
    V.row("hold", "＋持股：(Anthropic 估值 × 比例＋OpenAI 估值 × 比例) ×（1 − 流動性折價）", "US$bn",
          {"C": f"=(ant_val*{P('ant_sh')}+oai_val*oai_sh)*(1-{P('liq_disc')})"})
    V.row("eq", "股權價值", "US$bn", {"C": f"={v('ev', 'C')}-{v('nd26', 'C')}+{v('hold', 'C')}"})
    V.row("ps", "每股價值（2026-12-31；÷ 流通股數 sh_out：SBC 已視為現金成本，不另計稀釋）", "US$", {"C": f"={v('eq', 'C')}/sh_out"}, "0.00")
    V.row("dcf12", "DCF 情境價值推到目標價時點（×（1＋ke）^tp_yf − 期間股利）", "US$", {"C": f"={v('ps', 'C')}*(1+ke)^tp_yf-dps*tp_yf"}, "0.00", bold=True)
    V.sec("分部加總（SOTP）：目標價時點的未來 12 個月 EBITDA（FY2027 權重 1 − tp_yf、FY2028 權重 tp_yf）× 同業 NTM EV/EBITDA 中位數；EBITDA 用 SBC 加回口徑（與同業倍數同口徑，SBC 按營收比例分攤）；AWS AI 用自身 DCF（F3）")
    sbcw = lambda num, c: f"{P('sbc_ab')}*{k_('sbc', c)}*{num}/{S.t('tot', c)}"
    fw = lambda f: f"=(1-tp_yf)*({f('H')})+tp_yf*({f('I')})"
    for k, lab, mult, extra in (("ads", "廣告", f"m_ads*{P('m_x')}", lambda c: sbcw(S.t('rev_ad', c), c)),
                                ("rt", "零售本業（讀法 2：加回自用 AI 營運成本，自用 AI 價值中性）", f"m_retail*{P('m_x')}",
                                 lambda c: sbcw(f"({S.t('rt', c)}-{S.t('rev_ad', c)})", c) + f"+{rd}*(1-{P('ext_sh')})*{T.t('opex', c)}"),
                                ("awso", "AWS 其他（非 AI 雲端）", f"m_cloud*{P('m_x')}", lambda c: sbcw(S.t('aws_o', c), c))):
        V.row("e_" + k, f"{lab}：未來 12 個月 EBITDA（SBC 加回）", "US$bn", {"C": fw(lambda c, k=k, extra=extra: f"{l('e_' + k, c)}+{extra(c)}")})
        V.row("m_" + k, f"{lab}：倍數", "x", {"C": f"={mult}"}, "0.00")
        V.row("ev_" + k, f"{lab}：EV", "US$bn", {"C": f"={v('e_' + k, 'C')}*{v('m_' + k, 'C')}"})
    V.row("ai_fcf", "AWS AI 自由現金流＝AI EBITDA − 稅 × AI 營業利益 − 對外比例 × AI 資本支出（含汰換）", "US$bn",
          {c: f"={T.t('ebitda', c)}-tax*{T.t('ebit', c)}-{P('ext_sh')}*{T.t('cx', c)}" for c in FC})
    V.row("ai_pv", "  現值（WACC，期中）", "US$bn", {c: f"={v('ai_fcf', c)}*{v('df', c)}" for c in FC})
    V.row("ai_tv", "  永續價值（2032 年底）＝FY2033 AI NOPAT ×（1 − g ÷ RONIC）÷（WACC − g）", "US$bn",
          {"C": f"={T.t('ebit', 'M')}*(1-tax)*(1+{P('tg')})*(1-{P('tg')}/{P('ronic')})/({P('wacc')}-{P('tg')})"})
    V.row("ev_ai_d", "AWS AI：EV（自身 DCF，FY2028 起的現金流推到目標價時點；FY2027 現金流已在 2027 年底淨負債中）", "US$bn", {"C": f"=(SUM({v('ai_pv', 'I')}:{v('ai_pv', 'M')})+{v('ai_tv', 'C')}/(1+{P('wacc')})^6)*(1+{P('wacc')})^tp_yf"})
    V.row("e_ai", "對照：AWS AI 未來 12 個月 EBITDA", "US$bn", {"C": fw(lambda c: l('e_ai', c))})
    V.row("ev_ai_m", "對照：AWS AI EV（EBITDA × m_ai 倍數）", "US$bn", {"C": f"={v('e_ai', 'C')}*{P('m_ai')}"})
    V.row("ev_ai", "AWS AI：採用 EV（ai_leg＝1：自身 DCF；0：倍數）", "US$bn", {"C": f"=IF({P('ai_leg')}=1,{v('ev_ai_d', 'C')},{v('ev_ai_m', 'C')})"})
    V.row("sev", "分部 EV 合計", "US$bn", {"C": "=" + "+".join(v("ev_" + k, "C") for k in ("ads", "rt", "awso", "ai"))})
    V.row("nd27", "− 淨負債（2027 年底，含融資租賃與融資義務）", "US$bn", {"C": f"={k_('nd', 'H')}+fin_debt"})
    V.row("seq", "股權價值（＋持股）", "US$bn", {"C": f"={v('sev', 'C')}-{v('nd27', 'C')}+{v('hold', 'C')}"})
    V.row("sotp", "分部加總每股價值（目標價時點）", "US$", {"C": f"={v('seq', 'C')}/sh_val"}, "0.00", bold=True)
    V.row("tp", "情境目標價＝w_dcf × DCF＋（1 − w_dcf）× 分部加總", "US$", {"C": f"={P('w_dcf')}*{v('dcf12', 'C')}+(1-{P('w_dcf')})*{v('sotp', 'C')}"}, "0.00", bold=True, key_fill=True)
    V.row("sotp_aim", "對照：分部加總每股價值（AWS AI 用 EBITDA × 倍數）", "US$", {"C": f"=({v('seq', 'C')}-{v('ev_ai', 'C')}+{v('ev_ai_m', 'C')})/sh_val"}, "0.00")
    V.close()
    return {"S": S, "T": T, "L": L, "C": C, "V": V}


SH = {}
for nm_ in ("Thesis", "Seg", "PL", "CF", "Val"):
    SH[nm_] = wb.create_sheet(nm_)
CUR["PL"] = 12        # PL 頁最上方保留給跨情境核對列（segment_check 讀第一個同名列）
CUR["CF"] = 14        # CF 頁最上方保留給收支關卡摘要
CUR["Val"] = 90       # Val 頁最上方保留給第 0 節目標價與同業倍數表
ENG = {s: engine(s, s, SH) for s in SC}
E_ = {s: {k: b.id for k, b in ENG[s].items()} for s in SC}
tk_ = lambda s, blk, key, col: f"«{E_[s][blk]}.{key}.{col}»"

# ---------------- Thesis／Seg 頁首說明 ----------------
SH["Thesis"]["A1"] = ("Thesis — 命題模組：AWS AI＝對外 AI 在役 MW × 每 MW 年收入（Tokenomics v21.7 IF_HoldEcon × k；新增 Trainium＝GB300 錨 × IF_NonNV_Hold_Trn 0.61，v0.3.1）。"
                      "AI 資本支出（含自用）→ IT 6 年／機房 20 年折舊（次年起提列）→ 對外部分進 AWS 營業利益、自用部分進北美／國際營業利益。三情境 A／B／C 依序排列")
SH["Thesis"]["A2"] = "情境差異集中在命題驅動（In 頁 J 節）：A＝公司容量目標速度＋新產能 k 價格軸高；B＝資本支出指引隱含速度（C10，下半年視窗）＋基準 k；C＝與 B 同速度（資本支出照指引）＋新產能 k 價格軸低。非 AI 業務三情境相同"
SH["Seg"]["A1"] = ("Seg — 營收分部層：官方分部 North America、International、AWS；10-K 營收線 Online stores、Physical stores、Third-party seller services、Advertising services、"
                   "Subscription services、Other（零售各線合計＝北美＋國際）與 AWS。AWS＝AWS AI（命題子列，Thesis）＋AWS 其他。核對列（綠底）每情境每年度皆須為 0")
SH["PL"]["A1"] = "PL — 分部營收 → 分部營業利益 → 合併營業利益 → 業外、稅 → 淨利 → EPS。廣告利潤率 m_ad 驅動北美／國際利益率（廣告營收依零售地區比例分攤；零售本業利益率＝分部利益扣廣告利潤後的殘差）"
SH["PL"]["A2"] = "第 4–10 列：跨情境核對（三情境絕對值合計；tools/segment_check 讀本區）；實績欄為公司揭露值的核對"
SH["CF"]["A1"] = "CF — 營運現金（淨利＋分部折舊＋SBC＋其他）− 資本支出（AI＋AWS 其他＋零售）− 融資租賃本金 → 自由現金流；股利 0、回購 0（Amazon 無）；現金低於最低現金即舉債。FY2026＝9M 實績＋Q4 模型"
SH["Val"]["A1"] = "Val — 目標價約 12 個月後（2027-10-31）：DCF（FCFF、WACC；基準日 2026-12-31）50%＋分部加總（同業 NTM EV/EBITDA）50%；情境權重 A 25%／B 50%／C 25%；五級投資建議"

# PL 頁：跨情境核對（第一個同名列，segment_check 讀取）
PLS = SH["PL"]
PLS.cell(4, 1, "跨情境核對（三情境絕對值合計；應全為 0）").font = BOLD
hdr(PLS, 5, HEAD[:16])
for i, (key, lab) in enumerate((("chk_rt", "檢查：零售營收線合計 − (North America＋International)"), ("chk_aws", "檢查：AWS AI＋AWS 其他 − AWS"),
                                ("chk_ln", "檢查：營收線合計（含 AWS）− 合併營收"), ("chk_seg", "檢查：分部合計 − 合併營收"), ("chk_oi", "檢查：分部營業利益合計 − 合併營業利益")), 6):
    PLS.cell(i, 1, lab); PLS.cell(i, 2, "US$bn")
    for c in ALL:
        PLS[f"{c}{i}"] = "=" + "+".join(f"ABS({tk_(s, 'L', key, c)})" for s in SC); PLS[f"{c}{i}"].number_format = "0.000"
    for col in range(1, 16): PLS.cell(i, col).fill = CHK
    REG[("PLX", key)] = ("PL", i)

# CF 頁：收支關卡摘要
CFS = SH["CF"]
CFS.cell(4, 1, "收支關卡（地圖讀取；列標籤固定）").font = BOLD
hdr(CFS, 5, ["項目", "A（樂觀）", "B（基準）", "C（保守）", "情境加權", "說明"])
CFS.cell(6, 1, "收支關卡：自由現金流扣股利（五年合計）")
for j, s in enumerate(SC):
    CFS.cell(6, 2 + j, f"={tk_(s, 'C', 'gate', 'G')}").number_format = "#,##0.0"
CFS["E6"] = "=w_A*B6+w_B*C6+w_C*D6"; CFS["E6"].number_format = "#,##0.0"
CFS["F6"] = "FY2026–FY2030（首年起五年；FY2026＝9M 實績＋Q4 模型）自由現金流扣股利合計；負值＝入不敷出（地圖收支關卡）"
CFS.cell(7, 1, "FY2026–FY2030 新增借款合計（不含已發行債券）")
for j, s in enumerate(SC):
    CFS.cell(7, 2 + j, f"={tk_(s, 'C', 'borrow5', 'G')}").number_format = "#,##0.0"
CFS.cell(8, 1, "是否需要舉債")
for j, s in enumerate(SC):
    _c = openpyxl.utils.get_column_letter(2 + j)
    CFS.cell(8, 2 + j, f'=IF({_c}7>=1,"需要舉債",IF({_c}7>0,"幾乎不需要（新增借款少於 1）","不需要（現金足以支應）"))')
CFS["F8"] = "9M26 已發行債券（長期借款收入 115.1）在實績期內，不計；現金低於最低現金 50 時舉債"
CFS.cell(9, 1, "回購是否需停止")
for j in range(3):
    CFS.cell(9, 2 + j, "無回購可減（2023–9M26 皆為 0）")
CFS.cell(10, 1, "FY2028 自由現金流（扣股利）")
for j, s in enumerate(SC):
    CFS.cell(10, 2 + j, f"={tk_(s, 'C', 'fcfd', 'I')}").number_format = "#,##0.0"
for r_ in range(6, 11):
    for col in range(1, 7): CFS.cell(r_, col).fill = KEY if r_ == 6 else PatternFill()
REG[("CFX", "gate")] = ("CF", 6)
CFS.cell(11, 1, "檢查：收支關卡＝首年起五年自由現金流扣股利（逐欄加總；三情境絕對值合計）")
CFS["B11"] = "=" + "+".join(f"ABS({tk_(s, 'C', 'chk_gate', 'G')})" for s in SC); CFS["B11"].number_format = "0.000"
for col in range(1, 7): CFS.cell(11, col).fill = CHK

# ---------------- Val 第 0 節 ----------------
V = SH["Val"]
V.cell(3, 1, "0. 目標價（工作單第 2 節；AI 半導體 v2 口徑：DCF＋Comparable 各半，Comparable 在分部層做＝分部加總）").font = Font(bold=True, size=12)
hdr(V, 4, ["項目", "A（樂觀）", "B（基準）", "C（保守）", "說明"])
VR = {}
def vrow(r_, key, label, cells, fmt="#,##0.00", note=""):
    V.cell(r_, 1, label)
    for col, f in cells.items():
        V[f"{col}{r_}"] = f; V[f"{col}{r_}"].number_format = fmt
    if note: V.cell(r_, 5, note)
    VR[key] = r_; REG[("VX", key)] = ("Val", r_)
sc3 = lambda blk, key, col: {c: f"={tk_(s, blk, key, col)}" for c, s in zip("BCD", SC)}
vrow(5, "dcf12", "DCF 情境價值推到 12 個月後", sc3("V", "dcf12", "C"), note="FCFF 以 WACC 折到 2026-12-31，再以 ke 推到 2027-10-31")
vrow(6, "sotp", "Comparable（分部加總）情境價值", sc3("V", "sotp", "C"), note="廣告、零售本業、AWS 其他、AWS AI 各乘同業倍數；＋持股 − 淨負債")
vrow(7, "tps", "情境目標價（DCF 50%＋分部加總 50%）", sc3("V", "tp", "C"))
vrow(8, "tvsh", "永續價值占 DCF 企業價值", sc3("V", "tvsh", "C"), "0.0%")
vrow(9, "w", "情境權重", {"B": "=w_A", "C": "=w_B", "D": "=w_C"}, "0%")
vrow(10, "awsm28", "FY2028 AWS 營業利益率", sc3("L", "m_aws", "I"), "0.0%", "核心問題（一）")
vrow(11, "fcf28", "FY2028 自由現金流扣股利", sc3("C", "fcfd", "I"), "#,##0.0", "核心問題（二）")
vrow(12, "roic28", "FY2028 對外 AI ROIC", sc3("T", "roic", "I"), "0.0%", "對照 WACC（In 頁）")
vrow(13, "eps28", "FY2028 EPS", sc3("L", "eps", "I"), "0.00")
vrow(14, "gate", "收支關卡（各情境）：自由現金流扣股利 FY2026–FY2030 合計", sc3("C", "gate", "G"), "#,##0.0", "CF 頁第 6 列")
hdr(V, 16, ["項目", "數值", None, None, "說明"])
vrow(17, "dcftp", "DCF 目標價（情境加權）", {"B": "=w_A*B5+w_B*C5+w_C*D5"})
vrow(18, "comptp", "Comparable 目標價（市場主流同業法）", {"B": "=w_A*B6+w_B*C6+w_C*D6"}, note="分部加總（工作單：Comparable 在分部層做）")
vrow(19, "tpw", "目標價（DCF 與 Comparable 加權）", {"B": "=w_dcf*B17+(1-w_dcf)*B18"}, note="＝Σ 情境權重 × 情境目標價")
vrow(20, "px", "現價（2026-10-30 收盤）", {"B": "=px"}, note="Motley Fool／Yahoo Finance；MarketChameleon 同值")
vrow(21, "upw", "上檔空間：加權目標價", {"B": "=B19/px-1"}, "0.0%")
vrow(22, "upc", "上檔空間：Comparable 目標價", {"B": "=B18/px-1"}, "0.0%")
vrow(23, "tier", "上檔空間級距", {"B": '=IF(B21>=0.5,"大於 50%",IF(B21>=0.25,"25%–50%",IF(B21>=0,"0%–25%",IF(B21>-0.25,"0%～−25%",IF(B21>-0.5,"−25%～−50%","小於 −50%")))))'}, "@")
vrow(24, "rec", "投資建議", {"B": '=IF(B21>=0.5,"強力買進",IF(B21>=0.25,"買進",IF(B21>-0.25,"中立",IF(B21>-0.5,"放空","強力放空"))))'}, "@",
     "五級固定規則：≥ +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、≤ −50% 強力放空")
vrow(25, "cons", "市場共識目標價", {"B": "=cons_tp"}, note="Koyfin 60 位平均（2026-10-30，財報後）；MarketBeat 325.12")
vrow(26, "upcons", "上檔空間：市場共識目標價", {"B": "=cons_tp/px-1"}, "0.0%")
vrow(27, "pess", "悲觀情境價（C 情境 DCF）", {"B": "=D5"})
vrow(28, "band", "價格帶（DCF 不如預期～順利）", {"B": '=TEXT(D5,"#,##0")&"～"&TEXT(B5,"#,##0")'}, "@", "C 情境 DCF～A 情境 DCF（目標價時點）")
vrow(29, "gate1", "收支關卡（引用 CF 頁）：自由現金流扣股利五年合計（情境加權）", {"B": "=CF!E6"}, "#,##0.0")
vrow(30, "wacc", "WACC／ke／g", {"B": "=wacc", "C": "=ke", "D": "=tg"}, "0.00%")
V.cell(31, 1, "目標價方法")
V_R1 = 32
V["B31"] = ("DCF：FCFF＝NOPAT（評價口徑，含讀法 2 自用 AI 影子效益）＋分部折舊 − 資本支出＋其他營運項目（SBC 不加回、股數用流通股數；融資租賃負債列淨負債；另以融資租賃本金代理非現金取得的租賃資產扣除），WACC（CAPM：β Blume 調整、rf 2026-10-28、"
            "kd＝rf＋AA 利差、D 含類債務），FY2027–FY2032 期中折現，永續價值＝FY2033 NOPAT ×（1 − g ÷ RONIC）÷（WACC − g），RONIC＝WACC；2030–2032 新增 MW 收斂到 g；淨負債含融資租賃與融資義務；持股加回。"
            "分部加總：廣告、零售本業、AWS 其他＝未來 12 個月 EBITDA（SBC 加回口徑，與同業倍數一致）× 同業 NTM EV/EBITDA 中位數；AWS AI＝自身 DCF（EBITDA − 稅 − AI 資本支出含汰換；FY2028 起，FY2027 已在 2027 年底淨負債中）；減 2027 年底淨負債（含類債務）、加持股，÷ 稀釋股數。")
for k in ("tps",):
    for c in "BCD": V[f"{c}{VR[k]}"].fill = KEY
for k in ("tpw", "upw", "rec"):
    for c in "AB": V[f"{c}{VR[k]}"].fill = KEY
# 同業倍數表
RD1 = {}
P0 = 40
V.cell(P0, 1, "同業 NTM EV/EBITDA（StockAnalysis EV ÷ MarketScreener NTM EBITDA 共識，2026-10-07；沿用 v0.2.2 共用同業檔，本輪未更新）").font = BOLD
hdr(V, P0 + 1, ["同業", "代號", "NTM EV/EBITDA", "組別", "來源", "標記"])
PEERS = [("retail", "Walmart", "WMT", 17.72), ("retail", "Costco", "COST", 25.33), ("retail", "Target", "TGT", 8.77),
         ("ads", "Meta", "META", 10.79), ("ads", "Pinterest", "PINS", 6.91), ("ads", "The Trade Desk", "TTD", 6.53), ("ads", "AppLovin", "APP", 11.38),
         ("ads", "Reddit", "RDDT", 13.63), ("ads", "Snap", "SNAP", 7.43),
         ("cloud", "Oracle", "ORCL", 9.49), ("cloud", "IBM", "IBM", 12.51), ("cloud", "SAP", "SAP", 15.13), ("cloud", "Salesforce", "CRM", 10.81),
         ("cloud", "ServiceNow", "NOW", 20.63), ("cloud", "Akamai", "AKAM", 9.7), ("cloud", "Equinix", "EQIX", 21.2), ("cloud", "DigitalOcean", "DOCN", 24.49)]
GRP = {}
for i, (g, nm_, tkr, m) in enumerate(PEERS, P0 + 2):
    V.cell(i, 1, nm_); V.cell(i, 2, tkr); c = V.cell(i, 3, m); c.font = BLUE; c.number_format = "0.00"
    V.cell(i, 4, {"retail": "零售", "ads": "數位廣告（C18 共用組）", "cloud": "非 AI 雲端（C12 共用組）"}[g])
    V.cell(i, 5, {"retail": S_PEERR, "ads": S_PEERA, "cloud": S_PEERC}[g]); V.cell(i, 6, "Derived")
    GRP.setdefault(g, []).append(i)
mr = P0 + 2 + len(PEERS)
for j, (g, nmx, lab) in enumerate((("retail", "m_retail", "零售"), ("ads", "m_ads", "數位廣告"), ("cloud", "m_cloud", "非 AI 雲端"))):
    rr = mr + j
    V.cell(rr, 1, f"{lab} 中位數"); V.cell(rr, 3, f"=MEDIAN(C{GRP[g][0]}:C{GRP[g][-1]})").number_format = "0.00"; V.cell(rr, 4, f"{len(GRP[g])} 家")
    V.cell(rr, 1).font = BOLD
    defname(nmx, f"Val!$C${rr}")
V.cell(mr + 3, 1, "AWS AI：6×（In 頁 m_ai，對照表 D6）；對照 CRWV TTM 25.0×、ORCL TTM 16.5×、NBIS 260×（只作對照）")
V.column_dimensions["A"].width = 58
for c in "BCD": V.column_dimensions[c].width = 16
V.column_dimensions["E"].width = 70

# ---------------- Sens（v0.3.1：一次改一個 In 頁輸入〔敏感度低／高〕，整本以 LibreOffice 重算後的加權目標價；amazon/sens_v031.py 產生、建置時寫入數值） ----------------
SE = wb.create_sheet("Sens")
_INR = {row[0]: row for row in ROWS if row[0] != "sec"}
def _fmt(k, v):
    u = _INR[k][2]
    if isinstance(v, str): return v.lstrip("=")
    if u in ("%", "pt"): return f"{v * 100:.4g}%"
    return f"{v:,.4g}"
def _lh(k):
    _, _, _, lo, b, hi, *_ = _INR[k]
    return (lo if lo is not None and lo != b else None), (hi if hi is not None and hi != b else None)
def _desc(k):
    lo, hi = _lh(k); b = _INR[k][4]
    return f"{_fmt(k, b)} → " + "／".join(_fmt(k, x) for x in (lo, hi) if x is not None)
SCASES = []   # (案例鍵, 驅動, 設定說明, 鍵欄, {方向: [(In 鍵, 值)]})
def case(key, lab, note=None, keycol=None, sets=None):
    if sets is None:
        lo, hi = _lh(key); sets = {}
        if lo is not None: sets["低"] = [(key, lo)]
        if hi is not None: sets["高"] = [(key, hi)]
    SCASES.append((key, lab, note if note is not None else _desc(key), keycol or key, sets))
HE = ("tk_he_h100", "tk_he_gb200", "tk_he_gb300", "tk_he_vr200")
case("internal_ai_reading", "自用 AI 讀法 1（只計成本；C16）", "讀法 2 → 讀法 1", sets={"低": [("internal_ai_reading", 1)]})
case("he_all", "IF_HoldEcon 四世代同時取低／高（每 MW 年持有成本；k_長約的分母同步變動）",
     f"Tokenomics 低成本／高成本欄（tk_he_h100、tk_he_gb200、tk_he_gb300、tk_he_vr200 同時；GB300 {_INR['tk_he_gb300'][4]:.2f} → {_INR['tk_he_gb300'][3]:.2f}／{_INR['tk_he_gb300'][5]:.2f}）",
     keycol="he_all", sets={"低": [(k, _INR[k][3]) for k in HE], "高": [(k, _INR[k][5]) for k in HE]})
case("trn_hold", "Tokenomics IF_NonNV_Hold_Trn（cf_chip、cf_cx 同時跟隨）", _desc("trn_hold") + "（MW 速度沿用基準）")
case("cf_chip", "自研晶片係數改回舊 C2 口徑 1.0（收入與資本支出）", "0.61 → 1.0（同時改 cf_cx；MW 速度沿用基準）", keycol="cf_chip cf_cx", sets={"高": [("cf_chip", 1.0), ("cf_cx", 1.0)]})
case("trn_out", "Trainium 收入係數改用產出比 IF_NonNV_Out_Trn（資本支出仍用持有比）", "收入 0.61 → 0.5（基準產出比）／" + "／".join(f"{x:g}" for x in (TRO[0], TRO[2])),
     keycol="trn_out cf_chip", sets={"基準": [("cf_chip", TRO[1])], "低": [("cf_chip", TRO[0])], "高": [("cf_chip", TRO[2])]})
case("c10_win", "C10 校準視窗（基準＝下半年：Q4 公式＝（指引 − 上半年）÷ 2）", f"低＝只校準 Q4（速度 {PACE.get('pace_q4', 0):,.1f}）；高＝全年平均（指引 ÷ 4；速度 {PACE.get('pace_fy', 0):,.1f}）",
     keycol="pace_B", sets={"低": [("pace_B", PACE.get("pace_q4", 0))], "高": [("pace_B", PACE.get("pace_fy", 0))]})
case("c10_g", "2026 資本支出指引 ±10（約 220 → 210／230；C10 速度隨之重解）", f"速度 {PACE.get('pace_g210', 0):,.1f}／{PACE.get('pace_g230', 0):,.1f}",
     keycol="capex_g26 pace_B", sets={"低": [("capex_g26", 210), ("pace_B", PACE.get("pace_g210", 0))], "高": [("capex_g26", 230), ("pace_B", PACE.get("pace_g230", 0))]})
case("pace_x", "MW 速度乘數（三情境同乘）")
case("ext_sh", "對外比例（C23）", _desc("ext_sh") + "（總 MW 隨之變動）")
case("k_long", "k_長約改用長約下緣價格（新產能與期初機隊）")
case("k_spot", "k_現貨改用 AWS 牌價")
case("long_sh", "長約占比")
case("rpm_open", "期初 k 對上公司 AI run-rate（>25）；只作對照", "期初每 MW 收入＝ai_rr ÷ 期初 MW", sets={"高": [("rpm_open", "=ai_rr*1000/mw_open")]})
case("mw_ye25", "對外 AI 在役 MW（2025 年底）")
case("mw_open", "對外 AI 在役 MW（2026-06-30；MW 速度沿用基準）")
case("cust_sh", "新增 MW 中自研晶片比例")
case("colo_sh", "機房租用比例")
case("tk_life", "AI IT 折舊年限")
case("fac_life", "機房折舊年限")
case("cip_lag", "實績期 AI 在建延後轉入在役")
case("tk_pq", "Tokenomics P×Q 上限（只進上限檢查列）")
case("g_awso_lt", "AWS 其他長期成長")
case("cx_awso", "AWS 其他資本支出率")
case("awso_g0", "AWS 其他成長起點改用 9M26 年增", "FY2025 年增 → 9M26 年增")
for k, lab, _ in LINES:
    case(f"g_{k}", f"{lab} 長期成長")
case("m_ad", "廣告營業利益率")
case("step_na", "北美零售本業利益率每年改善")
case("step_in", "國際零售本業利益率每年改善")
case("oth_r", "其他營運項目 ÷ 營收（FY2027 起）")
case("lease_pr", "融資租賃本金償還")
case("sh_g", "稀釋股數年增")
case("min_cash", "最低現金")
case("inv_q4", "2026Q4 策略投資（Anthropic 額度）")
case("wacc_adj", "WACC", "−1pt／＋1pt")
case("beta", "β（Blume 調整後的單一來源）")
case("erp", "股權風險溢酬")
case("kd_spr", "AA 公司債利差（債務成本）")
case("tax", "稅率")
case("tg", "永續成長 g")
case("ronic", "RONIC＝WACC＋5pt", "RONIC＝WACC → WACC＋5pt")
case("m_x", "分部倍數乘數（零售、廣告、非 AI 雲端）")
case("ai_leg", "AWS AI 腿改用 EBITDA × 倍數（6×／4×／10×）", "預設 AI 自身 DCF（F3）→ EBITDA × m_ai", keycol="ai_leg m_ai",
     sets={"基準": [("ai_leg", 0)], "低": [("ai_leg", 0), ("m_ai", _INR["m_ai"][3])], "高": [("ai_leg", 0), ("m_ai", _INR["m_ai"][5])]})
case("sbc_ab", "分部加總計入未來授予的稀釋（EBITDA 不加回 SBC）", "加回 → 不加回", sets={"低": [("sbc_ab", 0)]})
case("ant_sh", "Anthropic 持股比例")
case("liq_disc", "持股流動性折價")
case("w_dcf", "DCF 權重")
SE["A1"] = ("Sens — 一次改一個輸入（In 頁敏感度低／高），整本活頁簿以 LibreOffice 重算後的加權目標價（Val 頁「目標價（DCF 與 Comparable 加權）」）與投資建議；"
            "由 amazon/sens_v031.py 產生，建置時寫入（數值）。另列情境 B 目標價、FY2028 AWS 營業利益率（B）與收支關卡")
SE["A2"] = ("相對基準＝加權目標價 ÷ 基準加權目標價 − 1；龍捲風依絕對值排序（HTML）。基準與各案例的其他輸入相同。"
            "v0.3.1 起依個股模型原則第 9 節：In 頁每個有低／高的鍵都有案例（F 欄「鍵」），IF_HoldEcon 四世代同時一案（he_all）；"
            "cons_tp、c_sprev26、c_speps26（市場共識）只作對照，不進模型（只被 Recon、README 與 Val 顯示格引用），故無案例。"
            "C10 的判斷替代（視窗、指引 ±10）另列兩案，速度依 calib_pace_v031.py 重解")
hdr(SE, 3, ["驅動", "設定", "方向", "加權目標價", "相對基準", "鍵", "投資建議", "情境 B 目標價", "FY2028 AWS 營業利益率（B）", "收支關卡 B（五年）", "收支關卡（加權）", "FY2027 資本支出（B）"])
r = 4
SE_ROWS = []; SE_JOBS = {}; SE_ROW_OF = {}
for key, lab, note, kc, sets in SCASES:
    for side in ("基準", "低", "高"):
        if side not in sets: continue
        SE.cell(r, 1, f"{lab}：{side}" if side != "基準" else lab); SE.cell(r, 2, note); SE.cell(r, 3, side); SE.cell(r, 6, kc)
        jk = f"{key}|{side}"
        SE_JOBS[jk] = [[k, v] for k, v in sets[side]]
        res = SENS.get(jk)
        if res is not None:
            c = SE.cell(r, 4, res["tp"]); c.number_format = "#,##0.00"; c.font = BLUE
            SE.cell(r, 5, f"=D{r}/Val!$B$19-1").number_format = "+0.0%;-0.0%;0.0%"
            SE.cell(r, 7, res["rec"])
            for cc, kk, fm in ((8, "tpB", "#,##0.00"), (9, "awsm28B", "0.0%"), (10, "gateB", "#,##0.0"), (11, "gateW", "#,##0.0"), (12, "cx27B", "#,##0.0")):
                c = SE.cell(r, cc, res[kk]); c.number_format = fm; c.font = BLUE
        SE_ROWS.append((key, side, r)); SE_ROW_OF[jk] = r; r += 1
SENS_LAST = r - 1
r += 1
SE.cell(r, 1, "基準加權目標價（Val 頁）"); SE.cell(r, 4, "=Val!B19").number_format = "#,##0.00"; SE.cell(r, 7, "=Val!B24")
SE.cell(r, 8, "=Val!C7").number_format = "#,##0.00"; SE.cell(r, 9, "=Val!C10").number_format = "0.0%"; SE.cell(r, 10, "=CF!C6").number_format = "#,##0.0"; SE.cell(r, 11, "=CF!E6").number_format = "#,##0.0"
SE_BASE = r; r += 1
SE.cell(r, 1, "敏感度重算時的基準值（應與上一列相同；不同表示敏感度已過期）")
if "base|base" in SENS:
    c = SE.cell(r, 4, SENS["base|base"]["tp"]); c.number_format = "#,##0.00"; c.font = BLUE
    SE.cell(r, 5, f"=D{r}-D{r - 1}").number_format = "0.000000"
SE_STALE = r; r += 2
SE.cell(r, 1, "敏感度「讀法 1 三情境」逐案明細（活公式；Val 第 0 節引用）").font = BOLD; r += 1
R0 = r
SE.cell(R0, 1, "自用 AI 讀法 1（只計成本）三情境（主值為讀法 2）").font = BOLD
hdr(SE, R0 + 1, ["情境", "情境價", "DCF", "分部加總"])
CUR["Sens"] = R0 + 7
SB = {"Seg": SE, "Thesis": SE, "PL": SE, "CF": SE, "Val": SE}
for i, s_ in enumerate(SC):
    e = engine(f"R1{s_}", s_, SB, {"internal_ai_reading": "1"}, tag="｜讀法 1")
    SE.cell(R0 + 2 + i, 1, SCN[s_])
    for j, k in enumerate(("tp", "dcf12", "sotp")):
        SE.cell(R0 + 2 + i, 2 + j, f"={e['V'].t(k, 'C')}").number_format = "0.00"
    RD1[s_] = R0 + 2 + i
SE.column_dimensions["A"].width = 46; SE.column_dimensions["B"].width = 40
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 模型 vs 市場共識、公司指引與 v0.3。只作事後對照，不為對上數字而調參數（共同規則第 4 節）"
hdr(RC, 3, ["項目", "對照值", "模型 A", "模型 B", "模型 C", "差距（B ÷ 對照 − 1）", "來源", "說明"])
rr_ = [4]
def rsec(t):
    RC.cell(rr_[0], 1, t).font = BOLD; rr_[0] += 1
def rec(label, ref, getter, srcs="", note="", fmt="#,##0.0", gap=True, check=False):
    r_ = rr_[0]
    RC.cell(r_, 1, label)
    if ref is not None:
        RC.cell(r_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        if getter is not None:
            RC.cell(r_, 3 + j, getter(s)).number_format = fmt
    if ref is not None and gap and getter is not None and not (isinstance(ref, str) and not ref.startswith("=")):
        RC.cell(r_, 6, f"=D{r_}/B{r_}-1").number_format = "0.0%"
    RC.cell(r_, 7, srcs); RC.cell(r_, 8, note)
    if check:
        for col in range(1, 9): RC.cell(r_, col).fill = CHK
    rr_[0] += 1
    return r_
YCOL = {2026: "G", 2027: "H", 2028: "I"}
rsec("A. 合併共識（MarketScreener 2026-10-08；S&P Global 2026-10-26；皆為財報前，財報後版本未取得）")
for y in (2026, 2027, 2028):
    c = YCOL[y]; yy = y % 100
    rec(f"FY{y} 營收", f"=c_rev{yy}", lambda s, c=c: f"={tk_(s, 'L', 'tot', c)}", "MarketScreener", "FY2026 模型含 9M 實績")
    if y < 2028:
        rec(f"FY{y} 營收（S&P）", f"=c_sprev{yy}", lambda s, c=c: f"={tk_(s, 'L', 'tot', c)}", "S&P Global（StockAnalysis，2026-10-26）")
    rec(f"FY{y} EBITDA", f"=c_ebitda{yy}", lambda s, c=c: f"={tk_(s, 'L', 'ebitda', c)}", "MarketScreener",
        "模型＝營業利益（財報口徑）＋分部 PP&E 折舊；差距約一半是折舊口徑（不含影音內容與營業租賃資產攤銷），約一半是 SBC（供應商 EBITDA 多為 SBC 加回）")
    rec(f"FY{y} 營業利益（EBIT）", f"=c_ebit{yy}", lambda s, c=c: f"={tk_(s, 'L', 'oi', c)}", "MarketScreener")
    rec(f"FY{y} 淨利（GAAP）", f"=c_ni{yy}", lambda s, c=c: f"={tk_(s, 'L', 'ni', c)}", "MarketScreener", "財報口徑；FY2026 含 9M Anthropic 評價利益")
    rec(f"FY{y} EPS（GAAP）vs 模型財報口徑 EPS", f"=c_eps{yy}", lambda s, c=c: f"={tk_(s, 'L', 'eps_fin', c)}", "MarketScreener", fmt="0.00")
    if y < 2028:
        rec(f"FY{y} EPS（調整後，S&P）vs 模型不含業外 EPS", f"=c_speps{yy}", lambda s, c=c: f"={tk_(s, 'L', 'eps_x', c)}", "S&P Global（2026-10-26）", "模型扣除業外其他（近似）", fmt="0.00")
    rec(f"FY{y} 資本支出", f"=c_capex{yy}", lambda s, c=c: f"={tk_(s, 'C', 'cx', c)}", "MarketScreener")
    rec(f"FY{y} 自由現金流", f"=c_fcf{yy}", lambda s, c=c: f"={tk_(s, 'C', 'fcf', c)}", "MarketScreener", "模型扣融資租賃本金；供應商口徑未揭露", gap=False)
rec("市場共識目標價（MarketBeat，第二來源）", "=cons_mb", None, "MarketBeat（2026-10-29）", "Koyfin 332.41 見 Val 頁", fmt="0.00", gap=False)
rsec("B. 分部與營收線共識（逐分部）")
for lab, key in (("North America", "na"), ("International", "intl"), ("AWS", "aws"), ("Advertising services", "rev_ad"), ("AWS AI（命題子列）", "aws_ai")):
    for y in (2026, 2027):
        rec(f"FY{y} {lab}", "未找到", lambda s, k=key, c=YCOL[y]: f"={tk_(s, 'L', k, c)}", "未找到：本輪未取得公開的分部共識（免費來源只有合併數）",
            "公司不揭露 AI 營收" if key == "aws_ai" else "", gap=False)
rsec("C. 公司指引與說法（驗證列；不用於校準，Q4 營收中點除外）")
rec("FY2026 現金資本支出 vs 指引約 220", "=capex_g26", lambda s: f"={tk_(s, 'C', 'cx', 'G')}", "2026Q3 新聞稿與法說", "Q4 26 AI 資本支出錨定指引（9M 實績 172.2，隱含 Q4 47.8）")
rec("2026Q4 營收指引中點 vs 模型 Q4 26", "=g26q4_mid", lambda s: f"={tk_(s, 'L', 'tot', 'F')}", "2026Q3 新聞稿", "指引 240–247；情境 B 以零售校準係數對齊中點（A／C 的 AWS 不同）")
rec("2026Q4 營業利益指引中點 vs 模型 Q4 26", "=(g26q4_oilo+g26q4_oihi)/2", lambda s: f"={tk_(s, 'L', 'oi', 'F')}", "2026Q3 新聞稿", "指引 28.5–33.5；只作驗證列（不校準）")
rec("C10 對帳（下半年視窗；情境 B 應為 0）", None, lambda s: f"={tk_(s, 'T', 'c10_gap', 'F')}", "Thesis 頁", "（指引 − 上半年實績）÷ 2 −（Q4 公式 AI＋非 AI）", gap=False)
rec("Q4 26 指引補足額（錨定值 − 公式值）", None, lambda s: f"={tk_(s, 'T', 'c10_q4', 'F')}", "Thesis 頁", "負值＝指引隱含的 Q4 支出低於 C10 速度的公式值（Q4 新增 MW 部分由已投入的在建支應）", gap=False)
rec("AWS 年增 FY2027 vs 公司「>40%」", "=aws_g40", lambda s: f"={tk_(s, 'L', 'aws', 'H')}/{tk_(s, 'L', 'aws', 'G')}-1", "2026Q3 新聞稿 CEO 引述、法說", "只作驗證列：不回頭改 MW、k 或 AWS 其他成長", fmt="0.0%", gap=False)
rec("AWS 年增 FY2028 vs 公司「>40%」", "=aws_g40", lambda s: f"={tk_(s, 'L', 'aws', 'I')}/{tk_(s, 'L', 'aws', 'H')}-1", "同上", "", fmt="0.0%", gap=False)
rec("AWS 年化營收 run-rate（Q4 26 × 4）vs 2026Q3 的 184", "=aws_rr", lambda s: f"={tk_(s, 'L', 'aws', 'F')}*4", "2026Q3 新聞稿", "")
rec("AWS backlog 578 vs 模型未來 24 個月 AWS 營收（Q4 26＋FY2027＋0.75 × FY2028）", "=aws_bl", lambda s: f"={tk_(s, 'L', 'aws', 'F')}+{tk_(s, 'L', 'aws', 'H')}+0.75*{tk_(s, 'L', 'aws', 'I')}", "2026Q3 法說",
    "backlog 認列期間未揭露（Q2 RPO 加權剩餘 6.4 年）；只作驗證列", gap=False)
rec("AWS AI 年化營收（2026-06-30）vs 公司「>25」", "=ai_rr", lambda s: "=mw_open*rpm_open/1000", "2026Q2 法說", "C3：不回頭改 MW 或 k，差距本身是發現")
rec("要對上公司說法：所需 k（MW 不變）", None, lambda s: "=ai_rr*1000/(mw_open*he_open)", "", "對照 k_B（In 頁）", fmt="0.000", gap=False)
rec("要對上公司說法：所需對外 MW（k 不變）", None, lambda s: "=ai_rr*1000/rpm_open", "", "對照期初 2,080 MW", fmt="#,##0", gap=False)
_rr = SE_ROW_OF.get("rpm_open|高")
rec("替代：期初 k 對上公司 run-rate 時的加權目標價變動（每股）", None, lambda s: f"=Sens!D{_rr}-Val!B19", "Sens 頁案例", "不作主值（共同規則第 4 節：不以公司說法反推 k）", fmt="0.00", gap=False)
rec("上限檢查：FY2028 每 MW 年收入 ÷ Tokenomics P×Q（VR200）", None, lambda s: f"={tk_(s, 'T', 'pq_chk', 'I')}", "Tokenomics v21.7", "上限，不得作預測", fmt="0.0%", gap=False)
rsec("D. v0.3 → v0.3.1：加權目標價變動拆解（(a)–(d) 為整本重算快照：Sens 頁基準差＝0 時有效；(e) 不影響目標價；(f) 為活公式）")
CHAIN = [("S1", "(a) 時間推移與期間結構：評價日 2026-06-30 → 09-30（FY2026 改 9M 實績＋Q4 模型）、目標日 2027-10-27 → 10-31；2026Q3 用 v0.3 自己的模型值", "v03"),
         ("S2", "(b) 新季實績：2026Q3 實績取代 v0.3 的模型值（分部營收與利益、營收線、利息、稅、淨利、現金流、資本支出、現金、債務、股數、策略投資）", "S1"),
         ("S3", "(c) 公司指引與 C10：Q4 26 營收指引 240–247（v0.3 隱含 Q4 改為指引中點）、2026 資本支出約 220 維持 → C10 速度重解（下半年視窗；舊 Trainium 口徑）", "S2"),
         ("S4", "(d) Tokenomics v21.4 → v21.7：Trainium 係數 1.0 → IF_NonNV_Hold_Trn 0.61（收入與 IT 資本支出；C10 速度依新係數重解）", "S3")]
RC_B0 = rec("v0.3 加權目標價（2026-10-27）", "=v03_tp", None, "v0.3 成品 Val!B19", fmt="0.00", gap=False)
_ok = f"ABS(Sens!E{SE_STALE})<0.000001"
CH_ROWS = []
for k, lab, prv in CHAIN:
    res = SENS.get(f"chain|{k}"); pv = SENS.get(f"chain|{prv}", {}).get("tp") if prv != "v03" else None
    f_ = (f"=IF({_ok},{res['tp']}-" + ("v03_tp" if prv == "v03" else f"{pv}") + ",\"快照過期\")") if res else None
    CH_ROWS.append(rec(lab, f_, None, f"重算快照 chain|{k}：目標價 {res['tp']:.4f}" if res else "（尚未重算）", "；".join(f"{a}={b}" for a, b in (res or {}).get("set", [])[:10]), fmt="0.00", gap=False))
s4 = SENS.get("chain|S4", {}).get("tp")
CH_ROWS.append(rec("(e) 現價與共識：現價 249.09 → 240.35（10-30）、共識 332.24 → 332.41；只影響上檔空間與評等", 0, None, "", "", fmt="0.00", gap=False))
CH_ROWS.append(rec("(f) 其他：無風險利率 5.00% → 4.99%、WACC 權益權重股價 249.09 → 240.35、Q4 其他營運項目改用 2025Q4 同期比例（C14）", f"=Val!B19-{s4}" if s4 else None, None, "活公式＝本版 − chain|S4", fmt="0.00", gap=False))
RC_SUM = rec("拆解合計（(a)–(f)）", "=" + "+".join(f"B{x}" for x in CH_ROWS) if SENS else None, None, "", fmt="0.00", gap=False)
RC_TOT = rec("總變動＝v0.3.1 − v0.3", f"=Val!B19-B{RC_B0}", None, "", fmt="0.00", gap=False)
rec("檢查：拆解合計 − 總變動（應為 0）", f"=B{RC_SUM}-B{RC_TOT}" if SENS else 0, None, "", fmt="0.000000", gap=False, check=True)
rec("另列：讀法 1 加權目標價 v0.3 → v0.3.1", "=v03_tp1", None, "", "", fmt="0.00", gap=False)
rec("另列：v0.3 上檔空間（當時現價 249.09）", "=v03_tp/v03_px-1", None, "", "v0.3 評等：中立", fmt="0.0%", gap=False)
RC.column_dimensions["A"].width = 64
for c in "BCDEF": RC.column_dimensions[c].width = 13
RC.column_dimensions["G"].width = 40; RC.column_dimensions["H"].width = 60
RECON_LAST = rr_[0] - 1

# Val 第 0 節：讀法 1 並列（F1）
V.cell(V_R1, 1, "讀法 1（自用 AI 只計成本）情境目標價")
for j, s_ in enumerate(SC):
    V.cell(V_R1, 2 + j, f"=Sens!B{RD1[s_]}").number_format = "0.00"
V.cell(V_R1, 5, "Sens 頁讀法 1 區塊；主值為讀法 2（In 頁 internal_ai_reading＝2）")
V.cell(V_R1 + 1, 1, "讀法 1：加權目標價"); V.cell(V_R1 + 1, 2, f"=w_A*B{V_R1}+w_B*C{V_R1}+w_C*D{V_R1}").number_format = "0.00"
V.cell(V_R1 + 2, 1, "讀法 1：上檔空間"); V.cell(V_R1 + 2, 2, f"=B{V_R1 + 1}/px-1").number_format = "0.0%"
V.cell(V_R1 + 3, 1, "讀法 1：投資建議")
V.cell(V_R1 + 3, 2, f'=IF(B{V_R1 + 2}>=0.5,"強力買進",IF(B{V_R1 + 2}>=0.25,"買進",IF(B{V_R1 + 2}>-0.25,"中立",IF(B{V_R1 + 2}>-0.5,"放空","強力放空"))))')
V.cell(V_R1 + 4, 1, "自用 AI 讀法（主值）"); V.cell(V_R1 + 4, 2, "=internal_ai_reading").number_format = "0"
V.cell(V_R1 + 4, 5, "2＝讀法 2（價值中性）；讀法 1 下 AI 容量越大目標價越低，讀法 2 下容量對目標價大致中性")

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources")
SO["A1"] = "Sources — 每個來源一列（In 頁、Q 頁以編號引用）。公司新聞稿與法說屬 Interested-party；10-K（經查核）與 10-Q（核閱）為 Verified；SemiAnalysis 未使用"
SO["A1"].font = BOLD
hdr(SO, 2, ["編號", "文件", "日期", "網址", "標記"])
for i, row in enumerate(SRC, 3):
    for j, v_ in enumerate(row, 1): SO.cell(i, j, v_)
SO.column_dimensions["A"].width = 8; SO.column_dimensions["B"].width = 90; SO.column_dimensions["C"].width = 12; SO.column_dimensions["D"].width = 100; SO.column_dimensions["E"].width = 18

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
B_ = lambda blk, key, col: tk_("B", blk, key, col)
lines = [
    ("Amazon（AMZN）個股模型 v0.3.1（命題模組＋官方營收分部 → 損益表；2026Q3 季度滾動）　2026-10-31", True),
    ("版本：v0.3.1（2026Q3 財報 2026-10-29 滾入；v0.3 為 2026-10-27）。資料截至 2026-09-30（2026Q3 新聞稿；10-Q 未取得，附註數字沿用 2026Q2）；FY2026＝9M 實績＋Q4 模型；Tokenomics v21.7（Trainium 改讀 IF_NonNV_Hold_Trn）。", False),
    ("命題：AWS 的 AI 容量（Trainium＋NVIDIA）能否以高於資金成本的報酬變現，零售與廣告的現金流能否支應資本支出。", True),
    ("核心問題：FY2028 AWS 營業利益率與自由現金流（扣股利後）是否轉正並維持。", True),
    ('="投資建議："&Val!B24&"（加權目標價 US$"&TEXT(Val!B19,"0.00")&"，現價 US$"&TEXT(px,"0.00")&"〔2026-10-30〕，上檔空間 "&TEXT(Val!B21,"+0.0%;-0.0%")&"；依五級固定規則）"', True),
    (f'="命題答案（情境 B）：FY2028 對外 AI ROIC "&TEXT({B_("T", "roic", "I")},"0.0%")&"，WACC "&TEXT(wacc,"0.0%")&"；打平所需機隊平均 k "&TEXT({B_("T", "k_be", "I")},"0.00")&"，模型 k "&TEXT({B_("T", "k_eff", "I")},"0.00")&"。FY2028 AWS 營業利益率 "&TEXT({B_("L", "m_aws", "I")},"0.0%")&"（FY2025 "&TEXT({B_("L", "m_aws", "D")},"0.0%")&"）；FY2028 自由現金流扣股利 "&TEXT({B_("C", "fcfd", "I")},"#,##0.0")&"；收支關卡（FY2026–FY2030 合計，情境加權）"&TEXT(CF!E6,"#,##0.0")&"。"', False),
    (f'="重點觀察 1：目標價 US$"&TEXT(Val!B19,"0.00")&"＝DCF "&TEXT(Val!B17,"0.00")&"（永續價值占 EV "&TEXT(Val!C8,"0%")&"）與分部加總 "&TEXT(Val!B18,"0.00")&" 各半；DCF 低是因為 WACC "&TEXT(wacc,"0.0%")&" 下，2027–2029 的 AI 資本支出（情境 B FY2027 "&TEXT({B_("C", "cx", "H")},"#,##0")&"）讓 FCFF 偏低，而 AI 增量報酬接近資金成本。"', False),
    (f'="重點觀察 2：分部加總中 AWS 其他（非 AI 雲端）占 "&TEXT({B_("V", "ev_awso", "C")}/{B_("V", "sev", "C")},"0%")&"、廣告占 "&TEXT({B_("V", "ev_ads", "C")}/{B_("V", "sev", "C")},"0%")&"；零售本業未來 12 個月 EBITDA（SBC 加回）"&TEXT({B_("V", "e_rt", "C")},"#,##0")&"；AWS AI 以自身 DCF 計 "&TEXT({B_("V", "ev_ai", "C")},"#,##0")&"。"', False),
    (f'="重點觀察 3：三情境目標價 A "&TEXT(Val!B7,"0")&"／B "&TEXT(Val!C7,"0")&"／C "&TEXT(Val!D7,"0")&"（差異來自新產能 k 與 MW 速度，In 頁 J 節）；讀法 1 加權 "&TEXT(Val!B33,"0.00")&"；市場共識目標價 "&TEXT(cons_tp,"0")&"；v0.3 加權目標價 "&TEXT(v03_tp,"0.00")&"（拆解見 Recon D 節）。"', False),
    ("", False),
    ("驅動 → 推導對照表（每頁輸入、推導量、因果方向）", True),
    ("In → Thesis：對外 AI MW（6/30 機隊 2,080 MW、評價日 9/30 加一季 C10 速度、之後每年新增 pace）× Tokenomics IF_HoldEcon × k → AWS AI 營收；新增 MW ÷ 對外比例 ×（IT＋自建機房成本）→ AI 資本支出 → IT 6 年／機房 20 年折舊 → AWS AI 營業利益（對外）與自用 AI 成本（零售）。", False),
    ("In → Seg：零售六條營收線（Q4 26＝Q4 25 ×（1＋9M26 年增）× 校準係數〔Q4 合併營收＝指引中點〕，之後收斂到長期年增）→ 零售合計 × 北美比例 → North America／International；AWS＝AWS AI＋AWS 其他（殘差起算、收斂到 8%）。", False),
    ("Seg、Thesis → PL：北美／國際營業利益＝廣告營收 × 50%＋零售本業營收 × 本業利益率（實績殘差＋每年改善）− 自用 AI 成本；AWS 營業利益＝AWS AI＋AWS 其他 × 近四季利益率 → 合併營業利益 → 利息（CF）、稅 17.4% → 淨利 → EPS。", False),
    ("PL → CF：淨利＋分部折舊＋SBC＋其他 − 資本支出（AI＋AWS 其他 20%＋零售 5.9%）− 融資租賃本金 → 自由現金流；股利 0、回購 0；現金低於 50 即舉債 → 利息費用回到 PL。", False),
    ("PL、CF → Val：DCF（FCFF、WACC、RONIC 永續）＋分部加總（廣告、零售本業、AWS 其他、AWS AI × 同業倍數）＋持股（Anthropic、OpenAI）→ 情境目標價 → 情境加權 → 上檔空間 → 投資建議。", False),
    ("", False),
    ("工作表：README → In（輸入；J 節情境驅動對照）→ Q（實績與核對）→ Thesis（命題模組）→ Seg（營收分部）→ PL → CF → Val → Recon（共識、指引、v0.3 拆解）→ Sens（敏感度）→ Sources。", False),
    ("標記：Verified、Interested-party、Analogy、Assumed、Derived。沿用 v0.2.2：k 規格（kspec.py）、對外比例、WACC 輸入、持股與同業倍數。SemiAnalysis 未使用。In 頁 D／F＝敏感度低／高，只用於 Sens 頁；共識區間只作對照。", False),
    ("已知限制：(1) 公司未揭露 AI 在役 MW、對外比例與 AI 營收，AWS AI／AWS 其他的歷史分拆為估計；(2) 分部共識未找到，合併共識為財報前；(3) 同業倍數沿用 2026-10-07；(4) 到期債務假設全數再融資；(5) 10-Q 未取得：債務面額、在建、分部折舊、租賃為推導或沿用 2026Q2。", False),
]
for i, (tx, b) in enumerate(lines, 1):
    RD.cell(i, 1, tx).font = Font(bold=b)
RD.column_dimensions["A"].width = 180

order = ["README", "In", "Q", "Thesis", "Seg", "PL", "CF", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    if ws.title in ("Thesis", "Seg", "PL", "CF"):
        ws.column_dimensions["A"].width = 58; ws.column_dimensions["P"].width = 60
        for c in "CDEFGHIJKLMNO": ws.column_dimensions[c].width = 11
        ws.freeze_panes = "C4"
resolve_all()
os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
wb.save(OUT)
ROWMAP = {"REG": {f"{a}|{b}": v for (a, b), v in REG.items()}, "INROW": INROW, "QR": QR, "VR": VR, "SENS_LAST": SENS_LAST, "SE_STALE": SE_STALE, "SE_BASE": SE_BASE,
          "SE_JOBS": SE_JOBS, "SE_ROWS": SE_ROWS, "RECON_LAST": RECON_LAST, "JROWS": JROWS,
          "CX27B": f"CF!H{REG[(E_['B']['C'], 'cx')][1]}", "Q4": {s_: {k: f"{sh}!F{REG[(E_[s_][b], k)][1]}" for b, k, sh in (("L", "tot", "PL"), ("L", "oi", "PL"), ("L", "aws", "PL"))} for s_ in SC}}
for _env in ("AMZN_ROWMAP", "ROWMAP"):
    if os.environ.get(_env):
        json.dump(ROWMAP, open(os.environ[_env], "w"), ensure_ascii=False, indent=1)


# =====================================================================================
# html.json（tools/company_html 展示層設定；列號取自本產生器，文字不含手寫數字）
# =====================================================================================
def html_config():
    rr = lambda s, blk, key: REG[(E_[s][blk], key)][1]
    ir = lambda k: INROW[k]
    hp = REG[(E_["B"]["L"], "hdr")][1]

    def drv(key, fmt, out_label, blk, okey, ocol, ofmt, direction):
        sh = {"T": "Thesis", "S": "Seg", "L": "PL", "C": "CF", "V": "Val"}[blk]
        return [f"@In!B{ir(key)}", [f"@In!E{ir(key)}|{fmt}", " ", f"@In!C{ir(key)}"], f"@In!G{ir(key)}|tag",
                [out_label, f"@{sh}!{ocol}{rr('B', blk, okey)}|{ofmt}"], direction]
    pl_metric = lambda name, key, fmt, chart=True, ymin=None, sheet="PL", blk="L": {
        "name": name, "fmt": fmt, "rows": {s: f"@{sheet}!G{rr(s, blk, key)}:M{rr(s, blk, key)}" for s in SC}, "chart": chart, **({"ymin": ymin} if ymin else {})}
    vfmt = {5: "n2", 6: "n2", 7: "n2", 8: "p0", 9: "p0", 10: "p1", 11: "n1", 12: "p1", 13: "n2", 14: "n1"}
    rec_fmt = {}
    for r_ in range(4, RECON_LAST + 1):
        lab = str(RC.cell(r_, 1).value or "")
        if "EPS" in lab or "目標價" in lab or "差異" in lab or "（1）" in lab or "（2）" in lab or "（3）" in lab:
            rec_fmt[str(r_)] = {c: "n2" for c in "BCDE"}
        if "所需 k" in lab:
            rec_fmt[str(r_)] = {c: "n3" for c in "CDE"}
        if "上限檢查" in lab:
            rec_fmt[str(r_)] = {c: "p1" for c in "CDE"}
        if lab.startswith("(") or "拆解" in lab or "總變動" in lab:
            rec_fmt[str(r_)] = {c: "n2" for c in "BCDE"}
        if "年增" in lab:
            rec_fmt[str(r_)] = {c: "p1" for c in "BCDE"}
        if lab.startswith("檢查"):
            rec_fmt[str(r_)] = {c: "n3" for c in "BCDE"}
        if "上檔空間" in lab:
            rec_fmt[str(r_)] = {"B": "sp1"}
    return {
        "company": "Amazon",
        "ticker": "AMZN",
        "thesis": "@README!A3",
        "question": "FY2028 AWS 營業利益率與自由現金流（扣股利後）是否轉正並維持",
        "scenarios": [{"key": "A", "name": "情境 A（樂觀）"}, {"key": "B", "name": "情境 B（基準）"}, {"key": "C", "name": "情境 C（保守）"}],
        "sections": [
            {"id": "thesis", "title": "命題與核心問題", "intro": "@README!A2", "blocks": [
                {"type": "kv", "items": [["結論", "@README!A5"], ["命題答案", "@README!A6"], ["三情境", "@Thesis!A2"], ["估值方法", "@Val!B31"]]},
                {"type": "h3", "text": "重點觀察"},
                {"type": "list", "items": ["@README!A7", "@README!A8", "@README!A9"]}]},
            {"id": "drivers", "title": "驅動因子 → 推導量", "intro": "標記沿用 Excel：Verified、Interested-party、Analogy、Assumed、Derived。情境驅動與單一參數區間分開呈現：前者三情境並排、進加權目標價；後者的敏感度低／高只用於敏感度頁。",
             "blocks": [
                {"type": "table", "caption": "情境驅動（三情境並排，進加權目標價；In 頁 J 節）", "head": [f"@In!B{JROWS[0] - 1}", f"@In!D{JROWS[0] - 1}", f"@In!E{JROWS[0] - 1}", f"@In!F{JROWS[0] - 1}", f"@In!C{JROWS[0] - 1}", f"@In!H{JROWS[0] - 1}"],
                 "rows": [[f"@In!B{jr}", f"@In!D{jr}|{fm}", f"@In!E{jr}|{fm}", f"@In!F{jr}|{fm}", f"@In!C{jr}", f"@In!H{jr}"] for jr, fm in zip(JROWS, ("n0", "n3", "p0"))]},
                {"type": "table", "caption": "單一參數區間（In 頁；基準進主模型，敏感度低／高只用於敏感度頁、不進目標價；三情境共用）", "head": ["@In!B3", "@In!D3", "@In!E3", "@In!F3", "@In!C3", "@In!G3"],
                 "rows": [[f"@In!B{ir(k)}", f"@In!D{ir(k)}|{fm}", f"@In!E{ir(k)}|{fm}", f"@In!F{ir(k)}|{fm}", f"@In!C{ir(k)}", f"@In!G{ir(k)}|tag"] for k, fm in (
                     ("mw_open", "n0"), ("pace_x", "n2"), ("ext_sh", "p0"), ("long_sh", "p0"), ("cust_sh", "p0"), ("trn_hold", "n2"), ("tk_he_gb300", "n2"), ("tk_life", "n0"),
                     ("g_awso_lt", "p0"), ("m_ad", "p0"), ("g_ad", "p0"), ("step_na", "p1"), ("erp", "p0"), ("tg", "p0"), ("ant_sh", "p0"), ("w_dcf", "p0"))]},
                {"type": "table", "caption": "驅動 → 推導量（情境 B）", "head": ["輸入", "基準值", "標記", "推導量（情境 B）", "因果方向"], "rows": [
                    drv("mw_open", "n0", "對外 AI 在役 MW（FY2028 年底）：", "T", "mw", "I", "n0", "六月底機隊是命題模組的起點；公司未揭露，取第三方估計"),
                    drv("pace_B", "n1", "新增 AI 總 MW（FY2027）：", "T", "addT", "H", "n0", "資本支出指引換成實體產能；除以對外比例得到含自用的總容量"),
                    drv("ext_sh", "p0", "自用 AI 成本淨額（FY2028）：", "T", "int_cost", "I", "n1", "對外比例越低，自用容量的成本越多落在零售分部"),
                    drv("k_B", "n3", "AWS AI 營收（FY2028）：", "T", "rev", "I", "n1", "每 MW 年收入＝持有成本 × k；k 由現貨與長約價格證據推算"),
                    drv("trn_hold", "n2", "新增 MW 加權 IT 資本支出（FY2028）：", "T", "cx_add", "I", "n2", "自研晶片的持有成本比同時降低每 MW 收入與每 MW 資本支出"),
                    drv("tk_cx_vr200", "n2", "AI 資本支出（FY2028）：", "T", "cx", "I", "n1", "新增總 MW 乘每 MW 的設備與自建機房成本，加上汰換"),
                    drv("tk_life", "n0", "AI 折舊（FY2028）：", "T", "da", "I", "n1", "資本支出依年限提列折舊，進入分部營業利益"),
                    drv("g_awso_lt", "p0", "AWS 其他營收（FY2028）：", "T", "awso", "I", "n1", "非人工智慧雲端由近期年增收斂到長期值"),
                    drv("m_ad", "p0", "北美營業利益（FY2028）：", "L", "oi_na", "I", "n1", "廣告利潤率決定零售分部利潤的來源；零售本業利益率為殘差"),
                    drv("g_ad", "p0", "廣告營收（FY2028）：", "S", "rev_ad", "I", "n1", "廣告是零售分部利益率的主要驅動"),
                    drv("wacc", "p1", "DCF 情境價值（情境 B）：", "V", "dcf12", "C", "n2", "折現率由資本資產定價模型計算"),
                ]},
                {"type": "h3", "text": "各頁因果"},
                {"type": "list", "items": ["@README!A12", "@README!A13", "@README!A14", "@README!A15", "@README!A16"]}]},
            {"id": "pl", "title": "三情境損益", "intro": "@PL!A1", "blocks": [
                {"type": "pl", "years": f"@PL!G{hp}:M{hp}", "metrics": [
                    pl_metric("合併營收（十億美元）", "tot", "n0"),
                    pl_metric("AWS 營收（十億美元）", "aws", "n0"),
                    pl_metric("其中 AWS AI（命題子列）", "aws_ai", "n0", chart=False),
                    pl_metric("AWS 營業利益率", "m_aws", "p1", ymin="auto"),
                    pl_metric("合併營業利益（十億美元）", "oi", "n0", chart=False),
                    pl_metric("每股盈餘（美元）", "eps", "n2"),
                    pl_metric("自由現金流（十億美元）", "fcf", "n0", sheet="CF", blk="C", ymin="auto"),
                    pl_metric("對外人工智慧投入資本報酬率", "roic", "p1", chart=False, sheet="Thesis", blk="T")]},
                {"type": "table", "caption": "分部核對（三情境絕對值合計，應為零）", "grid": True, "head": [f"@PL!A5", *[f"@PL!{c}5" for c in "CDEFGHIJKLM"]],
                 "rowgen": {"sheet": "PL", "rows": "6-10", "cols": ["A"] + list("CDEFGHIJKLM"), "fmts": [""] + ["n3"] * 11}}]},
            {"id": "valuation", "title": "估值", "intro": "@Val!A1", "blocks": [
                {"type": "table", "caption": "@Val!A3", "head": ["@Val!A4", "@Val!B4", "@Val!C4", "@Val!D4", "@Val!E4"],
                 "rowgen": {"sheet": "Val", "rows": "5-14", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "n2", "n2", "n2", ""],
                            "row_fmt": {str(k): {c: f for c in "BCD"} for k, f in vfmt.items()}}},
                {"type": "table", "head": ["@Val!A16", "@Val!B16", "@Val!E16"],
                 "rowgen": {"sheet": "Val", "rows": "17-30,33-36", "cols": ["A", "B", "E"], "fmts": ["", "n2", ""],
                            "row_fmt": {"21": {"B": "sp1"}, "22": {"B": "sp1"}, "26": {"B": "sp1"}, "29": {"B": "n1"}, "30": {"B": "p2"}, "34": {"B": "sp1"}}}},
                {"type": "bars", "title": "目標價與現價", "fmt": "n0", "items": [
                    {"label": "@Val!A19", "value": "@Val!B19", "series": 0}, {"label": "@Val!A17", "value": "@Val!B17", "series": 1},
                    {"label": "@Val!A18", "value": "@Val!B18", "series": 2}, {"label": "@Val!A25", "value": "@Val!B25", "series": 1}],
                 "ref": {"label": "@Val!A20", "value": "@Val!B20"}},
                {"type": "table", "caption": "收支關卡", "head": ["@CF!A5", "@CF!B5", "@CF!C5", "@CF!D5", "@CF!E5"],
                 "rowgen": {"sheet": "CF", "rows": "6-10", "cols": ["A", "B", "C", "D", "E"], "fmts": ["", "n1", "n1", "n1", "n1"]}},
                {"type": "note", "parts": ["@Val!A31", "：", "@Val!B31"]}]},
            {"id": "recon", "title": "對帳：共識、指引與前版", "intro": "@Recon!A1", "blocks": [
                {"type": "table", "head": [f"@Recon!{c}3" for c in "ABCDEFH"],
                 "rowgen": {"sheet": "Recon", "rows": f"4-{RECON_LAST}", "cols": list("ABCDEFH"), "fmts": ["", "n1", "n1", "n1", "n1", "sp1", ""], "row_fmt": rec_fmt}}]},
            {"id": "sens", "title": "敏感度", "intro": "@Sens!A1", "blocks": [
                {"type": "tornado", "title": "加權目標價相對基準的變動（一次改一個輸入，整本重算）", "sheet": "Sens", "rows": f"4-{SENS_LAST}", "label_col": "A", "value_col": "E", "fmt": "sp1"},
                {"type": "table", "collapse": True, "summary": "展開：全部案例的加權目標價、投資建議、情境 B 目標價、AWS 營業利益率與收支關卡",
                 "head": [f"@Sens!{c}3" for c in "ABDEGHIJK"],
                 "rowgen": {"sheet": "Sens", "rows": f"4-{SENS_LAST}", "cols": list("ABDEGHIJK"), "fmts": ["", "", "n2", "sp1", "", "n2", "p1", "n1", "n1"]}},
                {"type": "kv", "items": [[f"@Sens!A{SE_BASE}", f"@Sens!D{SE_BASE}|n2"], [f"@Sens!A{SE_STALE}", f"@Sens!D{SE_STALE}|n2"]]},
                {"type": "note", "parts": ["@Sens!A2"]}]},
            {"id": "status", "title": "資料狀態", "blocks": [
                {"type": "kv", "items": [["模型版本", ["$version", "（", "$date", "）"]], ["來源 xlsx", "$xlsx"], ["來源 xlsx 的雜湊指紋", "$sha"],
                                         ["Tokenomics 快照", [f"@In!H{ir('tk_he_h100')}"]]]},
                {"type": "h3", "text": "資料範圍"},
                {"type": "list", "items": ["@Q!A1", "@Q!A2", "@In!A2"]},
                {"type": "h3", "text": "已知限制"},
                {"type": "note", "parts": ["@README!A20"]},
                {"type": "h3", "text": "「未找到」清單"},
                {"type": "scan", "needle": "未找到", "collapse": True, "summary": "展開：標為「未找到」的項目", "sheets": {"Recon": ["A"], "In": ["B"]}, "empty": "沒有標為未找到的項目。"}]},
        ]}


if os.environ.get("AMZN_WRITE_HTMLJSON") or len(sys.argv) <= 1:
    json.dump(html_config(), open(os.path.join(HERE, "html.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(OUT)
