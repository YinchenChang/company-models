# v4.5（5a-1）：擷取「第一屏」文字——HTML 各頁面／分頁在 1440×900 視窗、捲動到頂端時可見的文字；
# Excel 各工作表前 40 個可見列（未被收合群組隱藏）的 A–J 欄文字（重算後的值）。
# 用途：已決定事項 9（第 3、4 輪）——滾動後，一頁摘要與各分頁第一屏的日期與期間文字不得過期；scripts/test_rolling.py 以此檢查。
# 用法：python3 scripts/first_screen.py 檔案.html|檔案.xlsx [輸出 json]
import sys, os, re, json

JS_VIEWPORT_TEXT = r"""() => {
  const out = [], H = window.innerHeight, W = window.innerWidth;
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  for (let n = w.nextNode(); n; n = w.nextNode()) {
    const t = n.textContent.replace(/\s+/g, ' ').trim(); if (!t) continue;
    const el = n.parentElement; if (!el) continue;
    const cs = getComputedStyle(el); if (cs.visibility === 'hidden' || cs.display === 'none') continue;
    const rg = document.createRange(); rg.selectNodeContents(n);
    const rs = [...rg.getClientRects()].filter(r => r.width > 0 && r.height > 0 && r.bottom > 0 && r.top < H && r.right > 0 && r.left < W);
    if (rs.length) out.push(t);
  }
  return out.join('\n');
}"""


def html_first_screens(path):
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from playwright.sync_api import sync_playwright
    out = {}
    with sync_playwright() as p:
        try: b = p.chromium.launch()
        except Exception: b = p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH', '/opt/pw-browsers/chromium'))
        pg = b.new_page(viewport={'width': 1440, 'height': 900})
        pg.goto('file://' + os.path.abspath(path)); pg.wait_for_timeout(2500)

        def snap(key):
            pg.evaluate('window.scrollTo(0, 0)'); pg.wait_for_timeout(150)
            out[key] = pg.evaluate(JS_VIEWPORT_TEXT)
        # 與 crawl.py 相同的導覽：三個頁籤、第一層選單、各自的次級分頁
        SKIP = ('總結', '資金模型', '損益與評價', '簡報模式（全螢幕）', '列印／另存 PDF', '載入', '重設', '儲存 JSON', '年度 CSV', '站點 CSV')
        def visible():
            return [t for t in dict.fromkeys(pg.locator('button:visible').all_inner_texts())
                    if 0 < len(t) < 14 and t not in SKIP and not t.startswith(('保守', '基準', '積極'))]
        def click(t):
            try:
                pg.locator('button:visible').filter(has_text=re.compile('^\\s*' + re.escape(t) + '\\s*$')).first.click(timeout=1500)
                pg.wait_for_timeout(400); return True
            except Exception: return False
        for top in ['總結', '資金模型', '損益與評價']:
            pg.get_by_role('button', name=top, exact=True).click(); pg.wait_for_timeout(700); snap(top)
            level1, states = visible(), {}
            for t in level1:
                if click(t): states[t] = visible(); snap(top + '/' + t)
            common = set.intersection(*(set(v) for v in states.values())) if states else set()
            for t, vis in states.items():
                for u in [u for u in vis if u not in common and top + '/' + u not in out]:
                    if click(t) and click(u): snap(top + '/' + t + '/' + u)
        b.close()
    return out


def xlsx_first_screens(path, rows=40):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb:
        lines, n = [], 0
        for r in range(1, ws.max_row + 1):
            if ws.row_dimensions[r].hidden: continue
            vals = [ws.cell(r, c).value for c in range(1, 11)]
            s = ' | '.join(str(v) for v in vals if v is not None and str(v).strip())
            if s: lines.append(s)
            n += 1
            if n >= rows: break
        out[ws.title] = '\n'.join(lines)
    return out


if __name__ == '__main__':
    f = sys.argv[1]
    o = xlsx_first_screens(f) if f.endswith('.xlsx') else html_first_screens(f)
    if len(sys.argv) > 2: json.dump(o, open(sys.argv[2], 'w'), ensure_ascii=False, indent=1)
    else:
        for k, v in o.items(): print(f'== {k}\n{v}\n')
