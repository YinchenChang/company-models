# 緯穎科技（6669）收支與估值模型 v0.1.1（AI 半導體 Project：伺服器組裝次產業，樣板＝quanta/）
# 用法：python3 wiwynn/build_wiwynn.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；整櫃價格與市場以美元輸入、以匯率 fx 換算；EPS：新台幣。
# 口徑：歸屬母公司淨利。業外與稅前利益以「淨利 ÷ 營業利益」一個係數涵蓋。
# 除權：2026-09-02 除權，每股配發約 1.98279460 股（盈餘轉增資＋可轉債轉換調整）。本模型所有 EPS 與股數皆為除權還原後
#       （FY2025、1H26 公司揭露 EPS ÷ 還原係數 2.98279460）；現價為除權後價格；年底 P/E 不受除權影響（價格與 EPS 同比例）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
QU_F = only("quanta/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ITC = None
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    if r[0].value == "AI IT 資本支出（加速器 × 每顆 IT capex）":
        ITC = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert ITC and all(v is not None for v in ITC[0]), ITC
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith("um") or str(k).startswith("u_") or str(k).startswith("g29_")):
        TS[k] = (r[4].value, r[0].row)
PLATS = ("bw", "rb", "ru", "hp", "amd", "tpu", "trn", "oth"); ASICP = ("trn", "oth")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
assert all(f"u_{p}_{y}" in TS for p in PLATS for y in (26, 27, 28)) and all(f"g29_{p}" in TS for p in PLATS), TS.keys()
def units(plats):
    u = {y: sum(TS[f"u_{p}_{y}"][0] for p in plats) for y in (26, 27, 28)}
    u[29] = sum(TS[f"u_{p}_28"][0] * (1 + TS[f"g29_{p}"][0]) for p in plats)
    return u
UTOT = units(PLATS); UASIC = units(ASICP)
U_SRC = f"{rel(TSMC_F)}｜In!E{TS['u_bw_26'][1]}:E{TS['g29_oth'][1]}"
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("racks", "capit"):
        BR[k] = (r[4].value, r[7].value, f"{rel(BR_F)}｜In!E{r[0].row}（GB300）、H{r[0].row}（VR200）")
assert len(BR) == 2, BR.keys()
# 廣達模型（組裝線權威，規則第 13 條）：ODM 每櫃價格、Rubin 占比、匯率、每 GW 價值變動；廣達有效份額只作 Recon
QWB = openpyxl.load_workbook(QU_F, data_only=True)
QU = {}
for r in QWB["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("odm_gb", "odm_vr", "vr26", "fx", "vpg_28", "vpg_29"):
        QU[k] = (r[3].value, r[4].value, r[5].value, f"{rel(QU_F)}｜In!D{r[0].row}:F{r[0].row}")
assert len(QU) == 6, QU.keys()
QSH = None
for r in QWB["Units"].iter_rows(min_row=4):
    if r[0].value == "廣達有效份額":
        QSH = ([r[2 + j].value for j in range(4)], f"{rel(QU_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert QSH and all(v is not None for v in QSH[0]), QSH

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B），整櫃價格與市場為美元；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個。EPS 與股數皆為除權還原後"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司原文＝緯穎財報與法說（經優分析、理財周刊、口袋學堂摘要）、月營收公告；公司的 AI 占比與展望為 Interested-party；券商與資料商估計皆為 Interested-party。組裝市場與 ODM 每櫃價格讀廣達模型（唯讀），口徑一致"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
KGI = "https://www.kgi.com.hk/zh-hk/-/media/files/kgishk/research-reports/tw-reports/wiwynn-6669-tt20260728en.pdf"
UAN2Q = "https://uanalyze.com.tw/articles/3734153022"
POCKET = "https://www.pocket.tw/school/articles/7853/"
BNEXT = "https://www.bnext.com.tw/article/92119/wiwynn-stock-dividend-1-to-3"
UAN0910 = "https://uanalyze.com.tw/articles/9272255032"
UAN1015 = "https://uanalyze.com.tw/articles/4794857151"
TN1008 = "https://finance.technews.tw/2026/10/08/wiwynn-6669-202609-financial-report/"
UANFEB = "https://uanalyze.com.tw/articles/4033848176"
MW0909 = "https://www.moneyweekly.com.tw/_Article?AID=247745"
CTEE0803 = "https://www.ctee.com.tw/news/20260803700886-430201"
UDN0921 = "https://money.udn.com/money/amp/story/5607/9768749"
CNY0615 = "https://news.cnyes.com/news/id/6498877"
SWS = "https://www.simplywall.st/stocks/tw/tech/twse-6669/wiwynn-shares/future"
INV = "https://www.investing.com/equities/wiwynn-consensus-estimates"
AKT = "https://aktien.guide/kursziel/Wiwynn-TW0006669005"
WEARN = "https://stock.wearn.com/cdata.asp?kind=6669"
SA = "https://stockanalysis.com/quote/tpe/6669/statistics/"
CMC = "https://companiesmarketcap.com/wiwynn/pe-ratio/"
GI = "https://goodinfo.tw/tw/StockBzPerformance.asp?STOCK_ID=6669&RPT_CAT=M_YEAR"
YEPS = "https://tw.stock.yahoo.com/quote/6669/eps"
URLNAME = {KGI: "凱基 2026-07-28 報告", UAN2Q: "優分析 2026-08-07（2Q26 財報）", POCKET: "口袋學堂（除權）", BNEXT: "數位時代（除權）",
           UAN0910: "優分析 2026-09-10（公司說明）", UAN1015: "優分析 2026-10-15", TN1008: "TechNews 2026-10-08（9 月營收）",
           UANFEB: "優分析 2026-02-26（4Q25 法說）", MW0909: "理財周刊 2026-09-09", CTEE0803: "工商時報 2026-08-03",
           UDN0921: "經濟日報 2026-09-21（高盛）", CNY0615: "鉅亨 2026-06-15（FactSet）", SWS: "Simply Wall St", INV: "Investing.com",
           AKT: "aktien.guide", WEARN: "聚財網（wearn）日線", SA: "stockanalysis.com", CMC: "companiesmarketcap", GI: "Goodinfo 經營績效", YEPS: "Yahoo 股市 EPS 頁"}

sec("A. AI 部署規模（讀 markets/network、markets/chip_bridge、tsmc 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, ITC[0][j], None, "Derived", ITC[1],
        "加速器顆數 × 每顆 IT capex（含 ASIC）；除以每 GW IT 資本支出得到部署 GW（與廣達、鴻海模型同口徑）" if y == 26 else "同上")
inp("itpg_gb", "每 GW IT 資本支出（GB300）", "US$B", None, BR["capit"][0], None, "Derived", BR["capit"][2], "Tokenomics v20.1 IF_CapexIT（經橋接表；唯讀）")
inp("itpg_vr", "每 GW IT 資本支出（VR200；2027 起）", "US$B", None, BR["capit"][1], None, "Derived", BR["capit"][2], "同上；ASIC 機架未另訂（Assumed：同 VR200）")
inp("vr26", "2026 年部署中 Rubin 世代占比（混合用）", "%", QU["vr26"][0], QU["vr26"][1], QU["vr26"][2], "Derived", QU["vr26"][3], "廣達模型輸入（TrendForce 2026 組合 Rubin 22%）")
inp("racks_gb", "每 GW 機架數（GB300）", "架", None, BR["racks"][0], None, "Derived", BR["racks"][2], "Tokenomics IF_RacksPerGW（經橋接表）")
inp("racks_vr", "每 GW 機架數（VR200）", "架", None, BR["racks"][1], None, "Derived", BR["racks"][2], "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", f"{rel(TSMC_F)}｜In!E{TS[f'um28_{s}'][1]}", "A＝順利（基準）；B＝平平（AI 資本支出消化）；C＝不如預期（明顯下修）；與台積電、廣達、鴻海模型相同")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", f"{rel(TSMC_F)}｜In!E{TS[f'um29_{s}'][1]}", "同上")
inp("fx", "匯率（新台幣 ÷ 美元）", "NT$/US$", QU["fx"][0], QU["fx"][1], QU["fx"][2], "Derived", QU["fx"][3], "與廣達、鴻海模型相同")
for y in (26, 27, 28, 29):
    inp(f"ua_{y}", f"緯穎客戶 ASIC 顆數 20{y}（AWS Trainium＋其他 ASIC）", "百萬顆", None, round(UASIC[y], 6), None, "Derived", U_SRC if y == 26 else "同上",
        "Trainium＋「其他 ASIC（MTIA、Maia、Ascend 等）」；Google TPU 不含（非緯穎主要客戶）；2029＝2028 × (1＋各類成長)" if y == 26 else "同上；「其他 ASIC」含華為 Ascend，非緯穎客戶（限制）")
for y in (26, 27, 28, 29):
    inp(f"ut_{y}", f"全部 AI 加速器顆數 20{y}", "百萬顆", None, round(UTOT[y], 6), None, "Derived", U_SRC if y == 26 else "同上", "八類加速器合計（與 Amkor 模型同口徑）" if y == 26 else "同上")

sec("B. 組裝價格（讀廣達模型，唯讀）、緯穎份額、ASIC 比重與客供")
inp("odm_gb", "ODM 每櫃價格 GB300（廣達模型）", "US$M/架", QU["odm_gb"][0], QU["odm_gb"][1], QU["odm_gb"][2], "Derived", QU["odm_gb"][3], "廣達模型標記 Interested-party（法人）；本模型不另訂；ASIC 機架每 GW 價值假設相同（Assumed）")
inp("odm_vr", "ODM 每櫃價格 VR200（廣達模型）", "US$M/架", QU["odm_vr"][0], QU["odm_vr"][1], QU["odm_vr"][2], "Derived", QU["odm_vr"][3], "同上")
inp("vpg_28", "每 GW ODM 整櫃價值 2028 ÷ 2027（廣達模型）", "x", QU["vpg_28"][0], QU["vpg_28"][1], QU["vpg_28"][2], "Assumed", QU["vpg_28"][3], "廣達模型的 Assumed")
inp("vpg_29", "每 GW ODM 整櫃價值 2029 ÷ 2027（廣達模型）", "x", QU["vpg_29"][0], QU["vpg_29"][1], QU["vpg_29"][2], "Assumed", QU["vpg_29"][3], "同上")
inp("asic_mix", "2026 年 ASIC 占緯穎 AI 伺服器營收", "%", 0.85, 0.90, 0.95, "Interested-party", UANFEB + "：ASIC 約占 AI 伺服器九成（公司，4Q25 法說）", KGI + "：2026F ASIC 91%、2027F 88%（凱基）")
inp("rel_27", "緯穎有效份額 2027 ÷ 2026（相對值，ASIC 傾斜之外）", "x", 0.85, 1.0, 1.1, "Assumed", UAN0910 + "：4Q26 ASIC 出貨超越通用伺服器、GPU 機櫃季底開始出貨（公司）", UDN0921 + "：AWS 新一代氣冷、液冷 ASIC 伺服器放量；NVIDIA、AMD 新 GPU 機櫃 4Q26 量產（高盛）")
inp("rel_28", "緯穎有效份額 2028 ÷ 2026", "x", 0.75, 1.0, 1.1, "Assumed", "同上；客戶自行分散 ODM（Celestica、廣達等）為下檔", "—")
inp("rel_29", "緯穎有效份額 2029 ÷ 2026", "x", 0.70, 0.95, 1.1, "Assumed", "同上", "—")
inp("cs_27", "客供比例 2027（AI 毛額價值中改為客戶供料、不計營收的比例）", "%", 0.0, 0.05, 0.15, "Assumed", UAN2Q + "：自 4 月起部分客戶記憶體改代採購，不計營收與成本（公司）", KGI + "：AMD Helios 機櫃的客供條件仍在協商（凱基）；量化：未找到")
inp("cs_28", "客供比例 2028", "%", 0.0, 0.10, 0.25, "Assumed", "同上；廣達模型 20%（GPU 機櫃為主）", "—")
inp("cs_29", "客供比例 2029", "%", 0.0, 0.15, 0.35, "Assumed", "同上；廣達模型 25%", "—")
inp("keep", "客供部分保留的毛利比例（相對買賣模式）", "x", 0.70, 0.90, 1.0, "Assumed", "公司：代採購使營收與成本同步減少（2Q26 財報）", "與廣達模型同值；量化：未找到")

sec("C. 營收結構錨點（公司揭露比例 × 營收）")
inp("ais_25", "2025 年 AI 伺服器 ÷ 營收", "%", 0.48, 0.50, 0.52, "Interested-party", UANFEB + "：2025 AI 約占五成，AI 與通用伺服器各半（公司）", KGI + "：2025 52%（凱基）")
inp("ais_26", "2026 年 AI 伺服器 ÷ 營收", "%", *SIL.vals("asm_wyai_26"), "Interested-party", UAN0910 + "：3Q26 ASIC 與通用伺服器出貨各半、4Q26 ASIC 超越（公司）；2026 年 AI 比重「超過 50%」（公司，2026-02）", KGI + "：2026F 46%（凱基 07-28；3Q 營收實績前）")
inp("g_nai27", "通用伺服器（非 AI）年增 2027", "%", -0.10, 0.05, 0.15, "Assumed", UDN0921 + "：北美 CSP 通用伺服器需求強、Agentic AI 帶動（高盛）", MW0909 + "：記憶體代採購使通用伺服器營收縮小（理財周刊）")
inp("g_nai28", "通用伺服器年增 2028", "%", -0.05, 0.03, 0.10, "Assumed", "同上", "—")
inp("g_nai29", "通用伺服器年增 2029", "%", -0.05, 0.03, 0.08, "Assumed", "同上", "—")

sec("D. 毛利率")
inp("d_ai", "AI 伺服器毛利率（毛額基礎）減通用伺服器毛利率", "pt", -0.02, -0.01, 0.0, "Assumed", UAN0910 + "：下半年 ASIC 比重上升，毛利率變化不大（公司）", KGI + "：毛利率 2026F 7.9% → 2027F 7.0%（AI 比重 46% → 61%）；部門別毛利率：未找到")
inp("vrd", "FY2027 起 AI 毛利率（毛額基礎）變動（GPU 機櫃比重上升）", "pt", -0.005, -0.002, 0.0, "Assumed", "GPU 機櫃含轉手 GPU，毛利率較低（公司：先前 GB200 專案拉低毛利率，4Q25 法說）", "廣達模型 VR 稀釋 −0.3 個百分點")

sec("E. FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設）")
inp("adj", "除權還原係數（1＋每股配股數）", "x", None, 2.9827946, None, "Verified", POCKET + "：2026-09-02 除權，每股配發 1.98279460 股，參考價 7,800 ÷ 2.9828 ＝ 2,615", BNEXT + "：每千股配 1,982.7946 股（可轉債轉換使配股率調整）")
inp("q3_rev", "3Q26 營收（月營收合計）", "NT$B", None, 434.618, None, "Verified", UAN1015 + "：3Q26 4,346.18 億、季增 56.25%；前三季 9,892.78 億", TN1008 + "：9 月 1,726.20 億、前 9 月 9,892.78 億（7 月 1,176.86、8 月 1,443.12，理財周刊）")
inp("q4g", "4Q26 營收季增", "%", -0.05, 0.08, 0.15, "Assumed", UAN1015 + "：4Q26 ASIC 出貨超越通用伺服器、GPU 機櫃季底出貨（公司）", "9 月單月 1,726 億 × 3 ＝ 5,179 億（季增 19%）；凱基 07-28 估 4Q26 季增 15%（以 3Q 估值為基）")
inp("q_gm", "3Q、4Q26 毛利率", "%", 0.076, 0.083, 0.088, "Interested-party", UAN0910 + "：下半年毛利率變化不大，NRE 貢獻待觀察（公司）；1H26 8.41%、2Q26 9.26%（含 NRE）", KGI + "：3Q26F 8.2%、4Q26F 7.4%")
inp("q_opx", "3Q、4Q26 營業費用 ÷ 營收", "%", 0.013, 0.015, 0.018, "Assumed", "1Q26 1.24%、2Q26 1.99%、1H26 1.62%（公司毛利率與營益率推得）", KGI + "：3Q26F 1.4%（8.2% − 6.8%）")
inp("q_nio", "3Q、4Q26 淨利 ÷ 營業利益", "%", 0.75, 0.78, 0.80, "Derived", "1H26 290.84 ÷ 376.8＝0.772；FY2025 511.18 ÷ 639.18＝0.800（公司）", "業外與稅前利益：未找到（摘要未列）")
inp("q_sh", "除權後股數（3Q、4Q26）", "十億股", None, 0.55736, None, "Interested-party", SA + "：流通股數 557.36 百萬股（2026-10-21）", "2Q26 加權股數 0.1861 × 還原係數 2.9828 ＝ 0.5551（Derived）")

sec("F. 營業費用、淨利係數、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.5, 0.8, 1.0, "Assumed", "FY2026 錨點實測見 Q 頁（約 1.0，含美國廠與研發擴編）", "廣達模型 0.4；營業費用率 FY2025 1.53%")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.002, 0.0, 0.002, "Assumed", "無", "營益率約 7%，±0.2 個百分點約 ±3% 的營業利益")
inp("nio_f", "淨利 ÷ 營業利益", "%", 0.72, 0.78, 0.80, "Derived", "同 q_nio（1H26 0.772、FY2025 0.800）", "低值：銀行聯貸 15 億美元與可轉債使利息費用上升（Assumed）")
inp("sh_g", "股數年變動", "%", 0.0, 0.005, 0.02, "Assumed", UAN2Q + "：董事會通過發行國內無擔保可轉債上限 150 億元（公司）", "150 億 ÷ 市值約 1 兆 ≈ 1.5%（轉換價：未找到）")
inp("mf", "AI 部署係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 AI IT 資本支出約 1.8 倍", "±20%")

sec("G. 券商、共識（只作 Recon 對帳，不作校準；除權前 EPS 已 ÷ 還原係數）")
inp("b_rev26", "凱基 FY2026 營收（07-28，3Q 前）", "NT$B", None, 1361.998, None, "Interested-party", KGI, "—")
inp("b_eps26", "凱基 FY2026 EPS（還原）", "NT$", None, "=370.69/adj", None, "Interested-party", KGI + "：370.69（除權前）", "—")
inp("b_rev27", "凱基 FY2027 營收", "NT$B", None, 2368.079, None, "Interested-party", KGI, "AI 伺服器 FY2027F 1,442（凱基）")
inp("b_eps27", "凱基 FY2027 EPS（還原）", "NT$", None, "=566.42/adj", None, "Interested-party", KGI + "：566.42（除權前）", "—")
inp("b_gm27", "凱基 FY2027 毛利率", "%", None, 0.070, None, "Interested-party", KGI, "營益率 5.8%")
inp("c_rev26", "共識 FY2026 營收（18 位）", "NT$B", None, 1403.066, None, "Interested-party", SWS + "（2026-10-22）", CNY0615 + "：FactSet 06-15 中位數 1,304（21 位）")
inp("c_rev27", "共識 FY2027 營收（19 位）", "NT$B", None, 2488.718, None, "Interested-party", SWS + "（2026-10-22）", CNY0615 + "：FactSet 06-15 中位數 1,943")
inp("c_ni26", "共識 FY2026 淨利（18 位）", "NT$B", None, 73.089, None, "Interested-party", SWS + "（2026-10-22）", "—")
inp("c_ni27", "共識 FY2027 淨利（19 位）", "NT$B", None, 109.350, None, "Interested-party", SWS + "（2026-10-22）", "—")
inp("c_ni28", "共識 FY2028 淨利（13 位）", "NT$B", None, 130.798, None, "Interested-party", SWS + "（2026-10-22）", "—")
inp("c_eps26", "共識 FY2026 EPS（還原；Simply Wall St 10-20）", "NT$", None, 138, None, "Interested-party", SWS + "：FY2026 EPS 估計由 123 上修至 138（2026-10-20）", "FactSet 09-02 中位數 361.05（除權前，18 位；數位時代）÷ 還原係數 ≈ 121")
inp("c_eps27", "共識 FY2027 EPS（還原；共識淨利 ÷ 除權後股數）", "NT$", None, "=c_ni27/q_sh", None, "Derived", SWS + " 共識淨利 ÷ 557.36 百萬股", "FactSet 06-15 中位數 463.51（除權前）÷ 還原係數 ≈ 155")
inp("c_eps28", "共識 FY2028 EPS（還原；共識淨利 ÷ 除權後股數）", "NT$", None, "=c_ni28/q_sh", None, "Derived", "同上", "FactSet 06-15 中位數 534.87（除權前）÷ 還原係數 ≈ 179")

sec("H. 價格與市場倍數")
inp("px", "股價（2026-10-22 收盤，除權後）", "NT$", None, 1840, None, "Interested-party", WEARN + "：10-22 收盤 1,840（開 1,940、低 1,830）", AKT + "：10-22 1,840；除權前 9-01 收盤 7,800、9-02 除權參考價 2,615")
inp("pe_ttm", "近 12 月 P/E（stockanalysis，10-21）", "x", None, 19.20, None, "Interested-party", SA, "companiesmarketcap 5.74（未還原除權，不採用）")
inp("fpe", "前瞻 P/E（stockanalysis，10-21 價格 1,880）", "x", None, 11.73, None, "Interested-party", SA, "—")
inp("cons_feps", "共識前瞻 12 月 EPS（10-21 價格 1,880 ÷ 前瞻 P/E）", "NT$", None, "=1880/fpe", None, "Derived", "stockanalysis 前瞻 P/E 11.73（2026-10-21）", "共識前瞻 EPS 的直接數字：未找到")
inp("cons_tp", "市場共識目標價（賣方平均，除權後）", "NT$", 2262.98, 2845.66, 4000, "Interested-party", INV + "：19 位平均 2,845.66、最高 4,000（高盛）、最低 2,262.98（2026-10-23 擷取）", AKT + "：20 位平均 3,060（10-22）；Simply Wall St 10-20：2,821")
REV10 = [317, 857, 1811, 1636, 1869, 1926, 2929, 2419, 3605, 9507]
NI10 = [3.11, 12.1, 55.8, 61.7, 86.1, 86.5, 142, 120, 228, 511]
inp("nim10", "十年平均淨利率（2016–2025，稅後淨利 ÷ 營收）", "%", None, "=AVERAGE(" + ",".join(f"{n}/{v}" for n, v in zip(NI10, REV10)) + ")", None, "Derived",
    GI + "：各年營收與稅後淨利（億元；2016 317／3.11 … 2025 9,507／511）", "凱基 07-28：FY2025 營收 950,663、淨利 51,118 百萬元（同）；2016–2017 為初期低利潤率年")

sec("I. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%；美、台、日、韓同一值）", "DCF 折現率；DCF 情境價值 ×(1＋r) 推到 12 個月後")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "DCF 永續成長率；FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均，FY2023–25）", "%", None, "=AVERAGE(0.297,0.349,0.48)", None, "Derived", GI + "：ROE FY2023 29.7%、FY2024 34.9%、FY2025 48%", "單一來源；stockanalysis 近 12 月 51.08%")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（Andy 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2029 淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（情境 B FY2029 與十年平均的中點）", "%", None, 0, None, "Derived", "PL 頁情境 B FY2029、In 頁 nim10", "—")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（＝十年平均淨利率）", "%", None, "=nim10", None, "Derived", "In 頁 nim10", "—")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    for u_, nm_ in URLNAME.items():          # 網址只列在 Sources 頁；In 頁寫來源名稱
        s1 = s1.replace(u_, nm_); s2 = s2.replace(u_, nm_)
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for k in TP_KEYS:
    for c in range(1, 10): IN.cell(INROW[k], c).fill = KEY
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣、除權還原後；股數為十億股、除權還原後）。藍字＝輸入；黑字＝公式。實績取自緯穎財報摘要與月營收"
Q["A2"] = "FY2026 錨點＝1H26 實績＋3Q26 營收實績（月營收）＋3Q、4Q 假設（毛利率、營業費用率、淨利 ÷ 營業利益；4Q 營收＝3Q × (1＋季增)）。2Q26 毛利與營業利益由公司揭露的毛利率、營益率 × 營收推得。EPS 還原：除權前 EPS ÷ 2.9828（2026-09-02 除權）"
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（營收實績）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：凱基 2026-07-28（公司年報數字）；1H26：優分析 2026-08-07（2Q26 財報）"
S2 = "Goodinfo：FY2025 營收 9,507 億、毛利率 8.25%、淨利 511 億、EPS 275.06"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "營收", {"B": 950.663, "C": 554.661, "D": "=q3_rev", "E": "=D4*(1+q4g)"}, tag="Verified／Assumed（4Q 季增）", s2="1Q26 2,765.08 億、2Q26 2,781.53 億（公司）")
qrow(5, "毛利", {"B": 78.454, "C": "=20.881+278.153*0.0926", "D": "=D4*q_gm", "E": "=E4*q_gm"}, tag="Verified／Derived（2Q26 毛利率 × 營收）", s2="1Q26 毛利 208.81 億；2Q26 毛利率 9.26%；1H26 毛利率 8.41%（公司）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(7, "營業費用", {"B": "=B5-B8", "C": "=C5-C8", "D": "=D4*q_opx", "E": "=E4*q_opx"}, tag="Derived（毛利 − 營業利益）／Assumed")
qrow(8, "營業利益", {"B": 63.918, "C": "=17.458+278.153*0.0727", "D": "=D5-D7", "E": "=E5-E7"}, tag="Verified／Derived", s2="1Q26 營業利益 174.58 億；2Q26 營益率 7.27%；1H26 營益率 6.79%（公司）")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(10, "歸屬母公司淨利", {"B": 51.118, "C": 29.084, "D": "=D8*q_nio", "E": "=E8*q_nio"}, tag="Verified／Derived（3Q、4Q）", s2="2Q26 149.69 億、1Q26 141.15 億（公司）")
qrow(11, "淨利 ÷ 營業利益（含業外、稅）", {c: f"={c}10/{c}8" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(12, "股數（十億股，除權還原後）", {"B": "=51.118/275.06*adj", "C": "=29.084/156.38*adj", "D": "=q_sh", "E": "=q_sh"}, ann="=(2*C12+D12+E12)/4",
     tag="Derived", s2="淨利 ÷ 除權前 EPS × 還原係數；FY2026＝半年加權（1H 兩季、3Q、4Q 各一季）")
qrow(13, "EPS（除權還原後）", {c: f"={c}10/{c}12" for c in "BCDE"}, fmt="0.00", ann="=F10/F12", tag="Derived", s1="淨利 ÷ 還原股數", s2="—")
qrow(14, "  公司揭露 EPS（除權前）", {"B": 275.06, "C": 156.38}, fmt="0.00", ann="", tag="Verified", s2="2Q26 80.43、1Q26 75.95（Yahoo 股市 EPS 頁同）")
qrow(15, "  公司揭露 EPS ÷ 還原係數", {"B": "=B14/adj", "C": "=C14/adj"}, fmt="0.00", ann="", tag="Derived", s1="—", s2="應等於第 13 列")
qrow(16, "AI 伺服器營收（＝營收 × AI 占比）", {"B": "=B4*ais_25"}, ann="=F4*ais_26", tag="Derived／Interested-party", s1="公司比例（In 頁 C 節）", s2="各季 AI 營收金額未揭露：未找到")
qrow(17, "通用伺服器與其他（非 AI）", {"B": "=B4-B16"}, ann="=F4-F16", tag="Derived", s1="—", s2="—")
Q.cell(19, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(20, 1, "FY2026 營收年增"); Q["B20"] = "=F4/B4-1"
Q.cell(21, 1, "FY2026 營業費用年增"); Q["B21"] = "=F7/B7-1"
Q.cell(22, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的對照）"); Q["B22"] = "=B21/B20"
Q.cell(23, 1, "FY2026 營收 ÷ 共識（Simply Wall St 10-22）− 1"); Q["B23"] = "=F4/c_rev26-1"
Q.cell(24, 1, "FY2026 EPS ÷ 共識（還原 138）− 1"); Q["B24"] = "=F13/c_eps26-1"
Q.cell(25, 1, "AI 伺服器營收 FY2026 年增（隱含）"); Q["B25"] = "=F16/B16-1"
Q.cell(26, 1, "1H26 EPS 自算 − 公司揭露還原（應約為 0）"); Q["B26"] = "=C13-C15"
for rr in range(20, 27): Q[f"B{rr}"].number_format = "0.0%" if rr not in (22, 26) else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 10), ("sh", 12), ("eps", 13), ("ai", 16), ("nai", 17)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$F${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 70; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I", 2032: "J"}
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
BU = (2026, 2027, 2028, 2029); F3 = (2027, 2028, 2029)
DEFP = dict(scn="A", rel="rel_{y}", cs="cs_{y}", keep="keep", vpg="vpg_{y}", mf="1", gn="g_nai{y}", d="d_ai", vrd="vrd", opxa="0",
            nio="nio_f", am="asic_mix", phi="phi", fg="fade_g{s}", fm="fade_m{s}")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    p = lambda k, y=None: P[k].format(y=(y % 100 if y else ""), s=s)
    um = lambda y: {2026: "", 2027: "", 2028: f"*um28_{s}", 2029: f"*um29_{s}"}[y]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 11, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 部署與整櫃組裝市場（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029）")
    row("itc", "AI IT 資本支出 × 情境顆數係數", "US$B", {c_(y): f"=itc_{y % 100}{um(y)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("itpg", "每 GW IT 資本支出", "US$B", {"D": "=itpg_gb*(1-vr26)+itpg_vr*vr26", **{c_(y): "=itpg_vr" for y in F3}})
    row("gw", "部署 GW（IT 電力）", "GW", {c_(y): f"={c_(y)}{R['itc']}/{c_(y)}{R['itpg']}" for y in BU}, fmt="0.0")
    row("vpg", "每 GW ODM 整櫃價值（廣達模型口徑）", "US$B/GW", {"D": "=(racks_gb*odm_gb*(1-vr26)+racks_vr*odm_vr*vr26)/1000", "E": "=racks_vr*odm_vr/1000",
                                                  **{c_(y): f"=$E{r}*{p('vpg', y)}" for y in (2028, 2029)}}, fmt="0.00",
        note="架數 × ODM 每櫃價格；ASIC 機架每 GW 價值假設相同（Assumed）")
    row("mku", "AI 伺服器組裝市場（ODM 價，美元）", "US$B", {c_(y): f"={c_(y)}{R['gw']}*{c_(y)}{R['vpg']}" for y in BU}, fmt="#,##0.0")
    row("mkn", "AI 伺服器組裝市場（新台幣）", "NT$B", {c_(y): f"={c_(y)}{R['mku']}*fx" for y in BU}, fmt="#,##0.0")
    row("sh0", "緯穎有效份額（FY2026 校準）＝ AI 伺服器營收 ÷ 市場", "%", {"D": f"=a26_ai/D{R['mkn']}"}, fmt="0.00%",
        note="Derived：分子為公司 AI 占比 × 營收；分母為全部 AI 組裝市場（含 NVIDIA、Google TPU 與其他組裝廠）")
    row("rel", "份額相對值（÷ FY2026；ASIC 傾斜之外）", "x", {"D": "=1", **{c_(y): f"={p('rel', y)}" for y in F3}}, fmt="0.00")
    row("tilt", "ASIC 傾斜係數 ＝ ASIC 比重 ×（客戶 ASIC 顆數指數 ÷ 全部顆數指數）＋（1 − ASIC 比重）", "x",
        {"D": "=1", **{c_(y): f"={P['am']}*((ua_{y % 100}/ua_26)/(ut_{y % 100}/ut_26))+(1-{P['am']})" for y in F3}}, fmt="0.000",
        note="Trainium＋其他 ASIC 成長快於全部加速器時，緯穎占 AI 組裝市場的份額上升；情境係數對兩者相同、互相抵銷")
    row("sh", "緯穎有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}*{c_(y)}{R['tilt']}" for y in BU}, fmt="0.00%")
    row("aiv", "AI 伺服器毛額價值（買賣模式下的營收）", "NT$B", {c_(y): f"={c_(y)}{R['mkn']}*{c_(y)}{R['sh']}" for y in BU}, fmt="#,##0.0")
    row("cs", "客供比例", "%", {"D": "=0", **{c_(y): f"={p('cs', y)}" for y in F3}}, fmt="0.0%",
        note="FY2026 的記憶體代採購已含在錨點與校準份額內，記為 0")
    row("ai", "AI 伺服器營收", "NT$B", {"C": "=a25_ai", **{c_(y): f"={c_(y)}{R['aiv']}*(1-{c_(y)}{R['cs']})" for y in BU}}, fmt="#,##0.0")
    sec_("B. 通用伺服器與合計")
    row("nai", "通用伺服器與其他（非 AI）", "NT$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{p('gn', y)})" for y in F3}}, fmt="#,##0.0")
    row("rev", "營收", "NT$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.0")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI 伺服器 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("pgw", "每 GW 緯穎 AI 伺服器營收", "US$B/GW", {c_(y): f"={c_(y)}{R['ai']}/fx/{c_(y)}{R['gw']}" for y in BU}, fmt="0.000")
    sec_("C. 損益")
    row("gmn", "通用伺服器毛利率 ＝ 錨點毛利率 − 差距 × AI 占比（FY2026 校準）", "%", {"D": f"=a26_gp/D{R['rev']}-{P['d']}*D{R['aish']}"}, fmt="0.00%",
        note="Derived：讓 FY2026 毛利率等於錨點；AI 毛利率（毛額基礎）＝通用 + 差距（差距為負）")
    row("gai", "AI 伺服器毛利 ＝ AI 毛利率 × 毛額價值 ×（1 − 客供 ×（1 − 保留比例））", "NT$B",
        {c_(y): f"=($D${R['gmn']}+{P['d']}" + ("" if y == 2026 else f"+{P['vrd']}") + f")*{c_(y)}{R['aiv']}*(1-{c_(y)}{R['cs']}*(1-{P['keep']}))" for y in BU}, fmt="#,##0.00",
        note="客供部分不計營收，但大部分毛利金額保留（keep）；FY2027 起加 GPU 機櫃稀釋 vrd")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"=$D${R['gmn']}*{c_(y)}{R['nai']}+{c_(y)}{R['gai']}" for y in F3}}, fmt="#,##0.0")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("gma", "AI 伺服器毛利率（營收基礎）", "%", {c_(y): f"={c_(y)}{R['gai']}/{c_(y)}{R['ai']}" for y in BU}, fmt="0.00%",
        note="客供比例上升 → 營收基礎毛利率上升")
    row("opb", "營業費用（基準路徑）", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,{P['phi']}*{c_(y)}{R['g']}))" for y in F3}}, fmt="#,##0.0")
    row("opx", "營業費用", "NT$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in F3}}, fmt="#,##0.0")
    row("oi", "營業利益", "NT$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.0")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("nio", "淨利 ÷ 營業利益", "%", {"C": "=a25_ni/a25_oi", "D": "=a26_ni/a26_oi", **{c_(y): f"={P['nio']}" for y in F3}}, fmt="0.0%")
    row("ni", "歸屬母公司淨利", "NT$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"={c_(y)}{R['oi']}*{c_(y)}{R['nio']}" for y in F3}}, fmt="#,##0.0")
    row("sh_", "股數（除權還原後）", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in F3}}, fmt="0.0000")
    row("eps", "EPS（除權還原後）", "NT$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("nim", "淨利率（歸屬母公司淨利 ÷ 營收）", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    # FY2030–32：收斂法（上游只到 2029；非驅動因子推導）
    for y in FADE:
        c, pc, k = FC[y], FC[y - 1], y - 2029
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+({p('fg')}-$G{R['g']})*{k}/3"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({p('fm')}-$G{R['nim']})*{k}/3"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['sh_']}"] = f"={pc}{R['sh_']}*(1+sh_g)"
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh_']}"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("g", "0.0%"), ("rev", "#,##0.0"), ("nim", "0.00%"), ("ni", "#,##0.0"), ("sh_", "0.0000"), ("eps", "0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fm; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 11, "FY2030 起為收斂法：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 緯穎損益（十億新台幣；EPS 新台幣、除權還原後；歸屬母公司）。FY2025A 實績、FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設；三情境相同），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2027 三情境相同（顆數係數從 CY2028 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數係數 1.0", "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 伺服器營收", "NT$B", "#,##0.0"),
        ("nai", "  通用伺服器與其他", "NT$B", "#,##0.0"), ("aish", "  AI 伺服器 ÷ 營收", "%", "0.0%"),
        ("gm", "毛利率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"), ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("nim", "  淨利率", "%", "0.00%"),
        ("eps", "EPS", "NT$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{SUMR[('A', 'nim')]}")
IN.cell(INROW["fade_mB"], 5, f"=(PL!G{SUMR[('B', 'nim')]}+nim10)/2")
for k in ("fade_mA", "fade_mB"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 錨點").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（NT$B）", "+".join(f"ABS(D{PB[s]['rev']}-a26_rev)" for s in SC)),
               ("  三情境：EPS 差（NT$）", "+".join(f"ABS(D{PB[s]['eps']}-a26_eps)" for s in SC)),
               ("  情境 A：毛利差（NT$B；以 FY2027 公式套用 FY2026）", f"D{PB['A']['gmn']}*D{PB['A']['nai']}+D{PB['A']['gai']}-a26_gp")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 60, sr
PL.cell(2, 1, "灰底＝FY2030–FY2032 收斂法（上游只到 2029）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 I 節），非驅動因子推導")
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["K"].width = 70
for c in "CDEFGHIJ": PL.column_dimensions[c].width = 13

# ---------------- Units ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — AI 伺服器組裝市場（讀廣達模型口徑）、緯穎份額、ASIC 傾斜與客供（情境 A）。部署 GW＝AI IT 資本支出（markets/network）÷ 每 GW IT 資本支出（Tokenomics，經橋接表）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("gw", "部署 GW（情境 A）", "GW", lambda y: f"=PL!{FC[y]}{A_['gw']}", "0.0", "AI IT 資本支出 ÷ 每 GW IT 資本支出"),
        ("mku", "AI 伺服器組裝市場（ODM 價）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}", "#,##0.0", "與廣達、鴻海模型同口徑"),
        ("ua", "緯穎客戶 ASIC 顆數（Trainium＋其他 ASIC）", "百萬顆", lambda y: f"=ua_{y % 100}", "0.00", "tsmc 成品"),
        ("ut", "全部 AI 加速器顆數", "百萬顆", lambda y: f"=ut_{y % 100}", "0.00", "tsmc 成品"),
        ("tilt", "ASIC 傾斜係數", "x", lambda y: f"=PL!{FC[y]}{A_['tilt']}", "0.000", "ASIC 比重 90% × 相對顆數指數＋10%"),
        ("sh", "緯穎有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.00%", "FY2026 以公司 AI 占比校準"),
        ("qsh", "對照：廣達有效份額（廣達模型）", "%", None, "0.00%", QSH[1]),
        ("aiv", "緯穎 AI 伺服器毛額價值", "US$B", lambda y: f"=PL!{FC[y]}{A_['aiv']}/fx", "#,##0.0", "買賣模式下的營收"),
        ("cs", "客供比例", "%", lambda y: f"=PL!{FC[y]}{A_['cs']}", "0.0%", "Assumed（FY2027 起）"),
        ("ai", "緯穎 AI 伺服器營收", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx", "#,##0.0", ""),
        ("pgw", "每 GW 緯穎 AI 伺服器營收", "US$B/GW", lambda y: f"=PL!{FC[y]}{A_['pgw']}", "0.000", ""),
        ("gma", "緯穎 AI 伺服器毛利率（營收基礎）", "%", lambda y: f"=PL!{FC[y]}{A_['gma']}", "0.00%", "客供上升 → 毛利率上升")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for j, y in enumerate(BU):
        v = f(y) if f else QSH[0][j]
        UN[f"{cyc(y)}{ur}"] = v; UN[f"{cyc(y)}{ur}"].number_format = fmt
        if not f: UN[f"{cyc(y)}{ur}"].font = BLUE
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：緯穎 AI 伺服器約九成是 ASIC 機架（AWS Trainium 為主，另有 Meta、Microsoft 自研晶片），所以份額隨「客戶 ASIC 顆數 ÷ 全部加速器顆數」上升；FY2026 有效份額只有約 3%，因為分母是全部 AI 組裝市場（含 NVIDIA 整櫃、Google TPU 與其他組裝廠）。份額相對值與客供比例是 Assumed。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 60; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣，除權後）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」；目標價＝DCF 與 Comparable 加權（第 0 節）；十年 P/E 規則值（第 3 節，年底落後 P/E）只作參考"
V["A2"] = "循環位置：緯穎不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照廣達、鴻海；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 51
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-22 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "#,##0"
for k, y, lab in ((3, 2026, "現價 ÷ FY2026 EPS（錨點）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 B）≈ 0.2 × FY2026＋0.8 × FY2027"); V[f"B{O + 7}"] = f"=0.2*PL!D{SUMR[('B', 'eps')]}+0.8*PL!E{SUMR[('B', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 B）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（stockanalysis，10-21）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis，10-21）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 3}"] = "FY2026 三情境相同（錨點）"; V[f"E{O + 4}"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, CY0 + 1, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
# 十年 P/E：companiesmarketcap 年底 P/E；Goodinfo 年底收盤 ÷ EPS（2019 年上市；2016–2018 無市價）。P/E 不受除權影響
PE_CMC = [None, None, None, 14.6, 12.3, 20.0, 8.99, 25.3, 19.6, 16.0]
PE_GI = [None, None, None, "=635/36.42", "=704/49.25", "=1115/49.46", "=797/81.07", "=1825/68.88", "=2620/126.57", "=4485/275.06"]
T10 = vr + 2
V.cell(T10, 1, "5. 十年 P/E（2016–2025 年底；兩來源平均，分位排名用；2019 年上市，2016–2018 無市價）").font = BOLD
hdr(V, T10 + 1, ["年底", "兩來源平均 P/E", "companiesmarketcap", "Goodinfo"])
RB0 = T10 + 16
for i, y in enumerate(range(2016, 2026)):
    rr = T10 + 2 + i
    V.cell(rr, 1, str(y))
    if y >= 2019:
        V.cell(rr, 2, f"=AVERAGE(C{rr}:D{rr})").number_format = "0.00"
        V.cell(rr, 3, f"=B{RB0 + 5 + i}").number_format = "0.00"
        V.cell(rr, 4, f"=C{RB0 + 5 + i}").number_format = "0.00"
PE10 = f"$B${T10 + 2}:$B${T10 + 11}"
V.cell(T10 + 12, 1, "年數"); V[f"B{T10 + 12}"] = f"=COUNT({PE10})"
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2019–2025 共 7 年；2017–2018 為上市前數字，不採用）"},
    {"label": "Goodinfo（年底收盤 ÷ EPS）", "vals": PE_GI, "tag": "Derived（資料商年底收盤與 EPS；除權前價格與 EPS，比值不受除權影響）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。緯穎 2019 年上市，只有 7 年（2019–2025）。兩來源相差 1–3 個倍數點（companiesmarketcap 可能用近 12 月 EPS）。2023 年 EPS 低谷時 P/E 25–26x，偏高。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025", (V.cell(RB0 + 5, 1).value, V.cell(RB0 + 14, 1).value)
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# 0. 目標價（Andy 2026-10-25 11:28／11:53：DCF 與 Comparable 各半；列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期）
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境（除權還原後）；FY2030 起為收斂法"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "0.00"); EPSR[y] = 8 + i
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
CFR = {}
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)" for c in "BCD"], "0.00"); CFR[y] = 16 + i
def dcf_formula(c, r="tp_r", g="tp_g", reinv="$B$15"):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], "0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("廣達（Quanta）", "2382.TW", 11.52, "2026-10-21"), ("緯創（Wistron）", "3231.TW", 9.21, "2026-10-19"),
         ("鴻海（Hon Hai）", "2317.TW", 11.94, "2026-10-21"), ("Super Micro Computer", "SMCI", 7.06, "2026-10-23"), ("Celestica", "CLS", 16.65, "2026-10-23")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "雲端伺服器 ODM／EMS 同業：廣達、緯創、鴻海、Supermicro、Celestica（不含 A 股）")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(2/12*B{EPSR[2027]}+10/12*B{EPSR[2028]})+tp_wB*(2/12*C{EPSR[2027]}+10/12*C{EPSR[2028]})+tp_wC*(2/12*D{EPSR[2027]}+10/12*D{EPSR[2028]})", None, None], "0.00",
     "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "", "", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁（0.5）")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "0.00", "Investing.com：19 位分析師平均 2,845.66（最高 4,000、最低 2,262.98），2026-10-23 擷取；aktien.guide 20 位 3,060")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-22 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝價格 ÷ 前瞻 P/E（In 頁 cons_feps）")
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
tset(46, "band", "價格帶（DCF 不如預期～順利）", ['=TEXT(D23,"0")&"～"&TEXT(B23,"0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2019–2025 年底 P/E（第 5 節兩來源平均）的分位；低於歷史最小值記 0%、高於最大值記 100%；只作檢查，不進目標價")
V.cell(48, 1, "目標價方法")
V["B48"] = ("DCF：股東現金流＝除權還原 EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為近三年平均約 37.5%（再投資率約 11%），FY2027–32 明確期＋永續；"
            "組裝廠資本支出低於淨利（2026 年約 13 億美元），未改用淨利＋折舊－資本支出。"
            "Comparable：廣達、緯創、鴻海、Supermicro、Celestica 前瞻 P/E 中位數（不含 A 股；英業達的 stockanalysis 價格停在 08-07，未列入）；"
            "Supermicro 倍數低（會計與毛利率疑慮）、Celestica 高（ASIC 與網通），中位數為廣達。"
            "偏誤：FY2030 起只有收斂法，淨利率高於十年平均時 DCF 偏高估；客戶集中（前三大約 96%）未在折現率反映。"
            "情境對應：A＝順利（顆數基準）、B＝平平（顆數 ×0.8／0.85）、C＝不如預期（顆數 ×0.6）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額、通用伺服器毛利率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("AI 部署 ×0.8", "FY2027 起 AI IT 資本支出 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("份額流失（低）", "份額相對值 2027 0.85、2028 0.75、2029 0.70（客戶分散 ODM）", dict(rel="rel_{y}_lo")),
         ("份額上升（高）", "份額相對值 1.1", dict(rel="rel_{y}_hi")),
         ("無 ASIC 傾斜", "緯穎份額只隨全部加速器成長（ASIC 比重設 0）", dict(am="0")),
         ("ASIC 比重 85%", "基準 90%", dict(am="asic_mix_lo")),
         ("客供 0（維持買賣）", "客供比例 2027–29 皆 0", dict(cs="cs_{y}_lo")), ("客供 高", "客供比例 15%／25%／35%", dict(cs="cs_{y}_hi")),
         ("客供保留毛利 0.7", "客供部分只保留 70% 毛利", dict(keep="keep_lo")),
         ("每 GW 價值 低", "2028–29 每 GW ODM 價值 ×0.85", dict(vpg="vpg_{y}_lo")), ("每 GW 價值 高", "×1.2", dict(vpg="vpg_{y}_hi")),
         ("通用伺服器成長 低", "2027 −10%、2028–29 −5%", dict(gn="g_nai{y}_lo")), ("通用伺服器成長 高", "2027 +15%、2028 +10%、2029 +8%", dict(gn="g_nai{y}_hi")),
         ("AI 毛利率差距 −2pt", "AI 毛利率（毛額基礎）比通用低 2 個百分點", dict(d="d_ai_lo")), ("AI 毛利率差距 0", "與通用相同", dict(d="d_ai_hi")),
         ("GPU 稀釋 −0.5pt", "FY2027 起 AI 毛利率 −0.5 個百分點", dict(vrd="vrd_lo")),
         ("營業費用率 +0.2pt", "營業費用額外 +0.2% × 營收", dict(opxa="opx_adj_hi")), ("營業費用彈性 1.0", "基準 0.8（FY2026 實測約 1.0）", dict(phi="phi_hi")),
         ("淨利 ÷ 營業利益 0.72", "FY2027 起 0.72（利息費用上升）", dict(nio="nio_f_lo"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "gm", 2027, "0.00%"), ("E", "eps", 2027, "0.00"),
                           ("G", "rev", 2028, "#,##0.0"), ("H", "gm", 2028, "0.00%"), ("I", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"F{i}"] = f"=E{i}/E$4-1"; SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c in "FJ": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 B", "CY2028 顆數 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "CY2028–29 顆數 ×0.6", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：顆數情境 B、C 的差異從 FY2028 起。緯穎的 AI 營收由部署 GW 與 ASIC 傾斜係數共同決定；「無 ASIC 傾斜」案例顯示上游 Trainium 與其他 ASIC 顆數假設對 EPS 的貢獻。客供只在保留毛利比例低時明顯影響 EPS。")
SENS_NOTE = nn + 3
TPS0 = nn + 5
SE.cell(TPS0, 1, "目標價參數（只改 DCF 參數；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "#,##0"
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.3", "", dict(w="0.3")), ("DCF 權重 0.7", "", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
assert TPS_LAST < 40, TPS_LAST
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 48
for c in "CDEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 券商估計、共識與份額口徑。只作事後對照，不為對上數字而調參數。除權前的券商 EPS 已 ÷ 還原係數 2.9828"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與券商（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=c_rev26", pl("rev", 2026), "Simply Wall St（18 位，10-22）", "", fmt="#,##0.0")
rec("FY2026 EPS vs 共識（還原）", "=c_eps26", pl("eps", 2026), "Simply Wall St（10-20 上修至 138）", "FactSet 09-02 中位數約 121（還原）")
rec("FY2027 營收 vs 共識", "=c_rev27", pl("rev", 2027), "Simply Wall St（19 位）", "", fmt="#,##0.0")
rec("FY2027 EPS vs 共識（共識淨利 ÷ 除權後股數）", "=c_eps27", pl("eps", 2027), "Simply Wall St（19 位）", "")
rec("FY2028 EPS vs 共識（共識淨利 ÷ 除權後股數）", "=c_eps28", pl("eps", 2028), "Simply Wall St（13 位）", "")
rec("FY2026 營收 vs 凱基（07-28，3Q 前）", "=b_rev26", pl("rev", 2026), "凱基 2026-07-28", "3Q26 營收 4,346 億高於凱基估 3,755 億", fmt="#,##0.0")
rec("FY2027 營收 vs 凱基", "=b_rev27", pl("rev", 2027), "凱基 2026-07-28", "", fmt="#,##0.0")
rec("FY2027 EPS vs 凱基（還原）", "=b_eps27", pl("eps", 2027), "凱基 2026-07-28", "")
rec("FY2027 毛利率 vs 凱基（百分點）", "=b_gm27", pl("gm", 2027), "凱基 2026-07-28", "", fmt="0.00%", pt=True)
rec("FY2027 AI 伺服器 ÷ 營收 vs 凱基 61%（百分點）", 0.61, pl("aish", 2027), "凱基 2026-07-28", "", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "B. 份額與每 GW").font = BOLD; rr_ += 1
rec("FY2026 緯穎有效份額 ÷ 廣達有效份額（廣達模型）", None, lambda s: f"=PL!D{PB[s]['sh0']}/Units!C{UR['qsh']}", "廣達模型 Units 頁", "同口徑（全部 AI 組裝市場）；緯穎約為廣達的四分之一", fmt="0.00")
rec("FY2026 每 GW 緯穎 AI 伺服器營收（US$B）", None, pl("pgw", 2026), "—", "", fmt="0.000")
rec("FY2027 ASIC 傾斜係數", None, pl("tilt", 2027), "tsmc 成品顆數", "大於 1：客戶 ASIC 顆數成長快於全部加速器", fmt="0.000")
rec("FY2027 AI 伺服器營收（NT$B）vs 凱基 1,442", 1442, pl("ai", 2027), "凱基 2026-07-28", "", fmt="#,##0.0")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("通用伺服器毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "AI 毛利率（毛額基礎）＝此值＋d_ai", fmt="0.00%")
rec("4Q26 EPS（模型，還原）", None, lambda s: "=Q!E13", "—", "", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代「AI IT 資本支出 ÷ 每 GW IT 資本支出」"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文只經財報摘要與新聞轉述，原始簡報待人工複核；券商與資料商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司（摘要）：優分析 2026-08-07（2Q26、1H26 營收、毛利率、營益率、淨利、EPS；記憶體代採購；資本支出 9.42 億美元）", UAN2Q),
        ("公司（摘要）：優分析 2026-09-10（3Q 通用與 ASIC 各半、4Q ASIC 超越、GPU 季底出貨、產能配置）", UAN0910),
        ("公司（摘要）：優分析 2026-02-26（4Q25 法說：AI 約五成、ASIC 約九成）", UANFEB),
        ("公司（摘要）：優分析 2026-10-15（3Q26 營收 4,346.18 億、前三季 9,892.78 億）", UAN1015),
        ("公司：月營收（TechNews 2026-10-08，9 月 1,726.20 億）", TN1008),
        ("公司（摘要）：理財周刊 2026-09-09（7、8 月營收；代採購說明）", MW0909),
        ("公司：除權（口袋學堂；配股率 1.98279460、參考價 2,615）", POCKET),
        ("公司：除權（數位時代；配股率調整原因、FactSet 2026 EPS 中位數 361.05）", BNEXT),
        ("券商：凱基 2026-07-28（FY2025 實績、2026F–2027F、客戶與 AI 占比、ASIC 比重）", KGI),
        ("券商：工商時報 2026-08-03（美系券商 EPS 2026–28 估計、目標價）", CTEE0803),
        ("券商：經濟日報 2026-09-21（高盛目標價 3,777、資本支出、ASIC 與 GPU 機櫃）", UDN0921),
        ("資料商：鉅亨 2026-06-15（FactSet 共識）", CNY0615),
        ("資料商：Simply Wall St（共識營收與淨利 2026–28、10-20 EPS 上修）", SWS),
        ("資料商：Investing.com（共識目標價 19 位）", INV),
        ("資料商：aktien.guide（共識目標價 20 位、10-22 價格）", AKT),
        ("資料商：聚財網日線（10-22 收盤 1,840）", WEARN),
        ("資料商：stockanalysis（股數、近 12 月與前瞻 P/E；同業前瞻 P/E）", SA),
        ("資料商：companiesmarketcap 年底 P/E", CMC),
        ("資料商：Goodinfo 歷年經營績效（年底收盤、EPS、ROE、營收、淨利）", GI),
        ("資料商：Yahoo 股市 EPS 頁（各季 EPS，除權前）", YEPS),
        ("市場層：markets/network、markets/chip_bridge、tsmc、quanta 成品（本 repo，唯讀）", f"{rel(NET_F)}；{rel(BR_F)}；{rel(TSMC_F)}；{rel(QU_F)}"),
        ("待補（人工）：3Q26 財報與法說（約 11 月上旬）；客戶占比（年報 10% 以上客戶，MOPS）；業外與稅前利益；ASIC 機架每櫃價格；客供比例的量化；可轉債轉換價；前瞻 P/E 歷史（Bloomberg）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("緯穎科技（6669）收支與估值模型 v0.1.1（AI 半導體 Project：伺服器組裝次產業，樣板＝quanta/）　2026-10-25", True),
 ("版本：v0.1.1（2026-10-25：補上檔空間級距、投資建議、重點觀察（數值不變））。v0.1（初版）。緯穎是雲端客戶（AWS、Microsoft、Meta）的伺服器 ODM，AI 伺服器約九成是 ASIC 機架（Trainium 為主），另一半營收是通用伺服器。組裝市場與 ODM 每櫃價格讀廣達模型（規則第 13 條），本模型只決定緯穎的份額、ASIC 傾斜、客供與通用伺服器。", False),
 ("模型命題：緯穎的 AI 伺服器營收＝部署 GW × 每 GW ODM 整櫃價值 × 緯穎有效份額（FY2026 校準）× ASIC 傾斜係數（客戶 ASIC 顆數相對全部加速器的成長）×（1 − 客供比例）；通用伺服器另以成長率推。毛利率約 8%，高於 GPU 整櫃組裝廠，因為 ASIC 機架不含轉手 GPU。", False),
 ("核心問題：客戶 ASIC（Trainium、MTIA、Maia）放量能不能讓 AI 營收在 FY2027 翻倍以上、同時守住約 8% 毛利率；除權後約 1,840 元、約 12 倍前瞻 P/E 隱含什麼路徑。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣；整櫃價格與市場以美元計、以匯率換算。口徑：歸屬母公司淨利。2026-09-02 除權（每股配約 1.98 股）：所有 EPS 與股數為除權還原後（除權前 EPS ÷ 2.9828），現價為除權後價格；年底 P/E 不受除權影響。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 伺服器營收：markets/network 的 AI IT 資本支出 × 情境顆數係數 ÷ 每 GW IT 資本支出（Tokenomics，經橋接表）＝部署 GW；× 每 GW 架數 × ODM 每櫃價格（廣達模型）＝組裝市場；× 緯穎有效份額（FY2026 以公司 AI 占比校準）× 份額相對值（Assumed）× ASIC 傾斜係數（tsmc 成品顆數，Derived）×（1 − 客供比例，Assumed）。", False),
 ("通用伺服器：FY2026 殘差 × 成長率（Assumed）。毛利率：AI 伺服器（毛額基礎）比通用低 d_ai（Assumed），FY2027 起再加 GPU 機櫃稀釋 vrd；通用毛利率為 FY2026 殘差（Derived）。營業費用彈性 phi（Assumed）。淨利：營業利益 × 淨利係數（業外與稅前未找到）。", False),
 ("校準：FY2026 ＝ 1H26 實績＋3Q26 營收實績＋假設（4Q 季增、毛利率、營業費用率）。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2）＋0.5 × Comparable 目標價（廣達、緯創、鴻海、Supermicro、Celestica 前瞻 P/E 中位數 × 12 個月後前瞻 EPS；不含 A 股）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解。十年 P/E 規則值 × FY2027 EPS 保留為參考（不是目標價）。情境對應：A＝順利、B＝平平、C＝不如預期（原為顆數情境）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文只經財報摘要與新聞；券商與資料商為 Interested-party。", False),
 ("", False),
 ("已知限制：(1) ASIC 傾斜係數用 tsmc 成品的「Trainium＋其他 ASIC」顆數，「其他 ASIC」含華為 Ascend（非緯穎客戶）；ASIC 機架每 GW 價值假設與 NVIDIA 整櫃相同；(2) 份額相對值與客供比例是 Assumed；(3) 客戶集中：前三大客戶約 96%（凱基，標籤順序不明），單一客戶轉單風險未建；(4) 業外與稅前利益未找到；(5) 十年 P/E 只有 2019–2025 七年；(6) 共識 FY2027–28 EPS 由共識淨利 ÷ 除權後股數推得；(7) 3Q26 財報（約 11 月上旬）後更新錨點；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ================= v0.1.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（工作單；Andy 2026-10-25 14:00、14:17） =================
# 全部放在 Val、README 現有最後一列之後，不位移既有儲存格
OBS = [
       "目標價 NT$3,314（DCF 3,756、Comparable 2,873 各半），比現價 1,840（2026-10-22）高 80.1%，級距大於 50%；Comparable 單獨為 56.1%。主因是 EPS 不是倍數：本模型 12 個月後前瞻 EPS 249.4 比共識前瞻 EPS 160.3 高 55.6%（兩者期間差約一年；FY2027 EPS 240.7 也比共識 196.2 高 22.7%），同業（廣達、緯創、鴻海、Supermicro、Celestica）前瞻 P/E 中位數 11.5 倍與現價隱含 11.5 倍相同；DCF 比 Comparable 高 30.7%，因 ROE 37.5% 使再投資率只有 10.7%（DCF 偏高估）。",
       "ASIC 傾斜是最大的單一假設：若緯穎份額只隨全部加速器成長（不向 ASIC 傾斜），FY2027 EPS −20.9%、FY2028 −23.4%。AWS Trainium 機架出貨與 3Q26 財報與法說（約 11 月上旬）是第一個檢驗點。",
       "情境分歧從 FY2028 起：FY2028 EPS A／B／C＝290.9／248.1／198.8，共識 234.7 介於 B 與 C 之間；份額相對值 0.75（客戶分散 ODM）時 FY2028 EPS −18.2%。"
      ]   # 重點觀察（與 README.md 相同文字；數字取自重算後的 xlsx）
_U = f"B{T['upw']}"
GR = V.max_row + 3
V.cell(GR - 1, 1, "0.4 上檔空間級距與投資建議（v0.1.1 新增；放在本頁最後以免位移既有儲存格；依「上檔空間：加權目標價」第 " + str(T["upw"]) + " 列）").font = BOLD
tset(GR, "grade", "上檔空間級距", [f'=IF({_U}>0.5,"大於 50%",IF({_U}>0.25,"25%–50%",IF({_U}>=0,"0%–25%",IF({_U}>=-0.25,"0%～−25%",IF({_U}>=-0.5,"−25%～−50%","小於 −50%")))))', None, None], "General",
     "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）；負值區間同樣重要")
tset(GR + 1, "rec", "投資建議", [f'=IF({_U}>0.5,"強力買進",IF({_U}>0.25,"買進",IF({_U}>=-0.25,"中立",IF({_U}>=-0.5,"放空","強力放空"))))', None, None], "General",
     "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）")
rn = RD.max_row + 1
RD.cell(rn, 1, f'="投資建議："&Val!B{GR + 1}&"（依上檔空間固定規則）"').font = BOLD
OBS_ROW = rn + 1
for i_, t_ in enumerate(OBS, 1):
    RD.cell(rn + i_, 1, f"重點觀察 {i_}：" + t_)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("WIWYNN_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE, "UN_NOTE": UN_NOTE,
               "T": T, "T10": T10, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "PEER_LAST": PEER_LAST, "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0, "NCASES": len(cases)},
              open(os.environ["WIWYNN_ROWMAP"], "w"), ensure_ascii=False, indent=1)
