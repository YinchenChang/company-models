# 智邦科技（2345）收支與估值模型 v0.2.1（AI 半導體 Project：網通系統端，白牌交換器 ODM）
# 用法：python3 accton/build_accton.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；市場以美元輸入、以匯率 fx 換算；EPS：新台幣。
# 口徑：稅後淨利（歸屬母公司；公司公告）。業外與稅以「淨利 ÷ 營業利益」一個係數涵蓋（比照 quanta/）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))   # 本資料夾的 tp_v2.py
import cycle_position as CP
import tp_v2 as TP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
wsn = openpyxl.load_workbook(NET_F, data_only=True)
NETR = {}
for r in wsn["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in ("scale-out 交換器", "AI IT 資本支出（加速器 × 每顆 IT capex）"):
        NETR[lab] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETR) == 2 and all(v is not None for v in NETR["scale-out 交換器"][0]), NETR.keys()
NU = {}   # 加速器顆數（百萬顆）：Units 頁
for r in wsn["Units"].iter_rows(min_row=4):
    lab = r[0].value
    for k, pat in (("trn", "AWS Trainium"), ("amd", "AMD Instinct"), ("asic", "其他 ASIC")):
        if lab and str(lab).startswith(pat):
            NU[k] = ([r[1 + j].value for j in range(4)], f"Units!B{r[0].row}:E{r[0].row}")
assert len(NU) == 3, NU.keys()
NI_ = {}  # 每顆 scale-up 內容（美元）：In 頁 E 欄
for r in wsn["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("su_cu", "c_TRN_su_sw", "c_TRN_su_ic", "c_ASIC_su_sw", "c_ASIC_su_ic", "c_AMD8_su_sw", "c_AMD8_su_ic", "c_AMDR_su_sw", "c_AMDR_su_ic"):
        NI_[k] = (r[4].value, f"In!E{r[0].row}")
assert len(NI_) == 9, NI_.keys()
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()

# 非 NVIDIA scale-up 乙太網路交換器（交換匣）市場：顆數 × （交換晶片＋交換匣其他 ×（1 − 銅纜占比））；TPU 為 ICI／OCS，不計
cu = NI_["su_cu"][0]
def tray(p): return NI_[f"c_{p}_su_sw"][0] + NI_[f"c_{p}_su_ic"][0] * (1 - cu)
SU = {}
for j, y in enumerate((26, 27, 28, 29)):
    amdp = "AMD8" if y == 26 else "AMDR"
    terms = [(NU["trn"][0][j], "TRN"), (NU["asic"][0][j], "ASIC"), (NU["amd"][0][j], amdp)]
    f = "=(" + "+".join(f"{u}*({NI_[f'c_{p}_su_sw'][0]}+{NI_[f'c_{p}_su_ic'][0]}*(1-{cu}))" for u, p in terms) + ")/1000"
    SU[y] = f
SU_SRC = f"{rel(NET_F)}｜{NU['trn'][1]}（Trainium）、{NU['asic'][1]}（其他 ASIC）、{NU['amd'][1]}（AMD）× In 頁 c_TRN／c_ASIC／c_AMD8（2026）／c_AMDR（2027 起）_su_sw、_su_ic 與 su_cu"

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B），市場為美元；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝智邦季報與月營收公告（經經濟日報、工商時報、富果摘要轉述）；法人估計與券商目標價皆為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
UDN_Q2 = "https://money.udn.com/money/story/5710/9675501"
UDN_Q1 = "https://money.udn.com/money/story/5710/9488861"
CT_Q3 = "https://www.chinatimes.com/realtimenews/20261006003734-260410"
CT_0929 = "https://www.ctee.com.tw/news/20260929700959-430201"
FUGLE = "https://blog.fugle.tw/post/earnings-call-2345-2025-11-7"
GI = "https://goodinfo.tw/tw/StockBzPerformance.asp?STOCK_ID=2345&RPT_CAT=M_YEAR"
SINO = "https://aiinvest.sinotrade.com.tw/Stock/Content/TW/2345"
FS = "https://news.cnyes.com/news/id/6629903"
UDN_SEO = "https://udn.com/news/story/7252/9832440"

sec("A. AI 交換器市場（讀 markets/network 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mkt_{y}", f"AI 後端 scale-out 交換器市場 20{y}", "US$B", None, NETR["scale-out 交換器"][0][j], None, "Derived", NETR["scale-out 交換器"][1],
        "加速器顆數 × 每顆 scale-out 交換器內容（不分乙太網路與 InfiniBand）；800G → 1.6T 的單價上升已含在每顆內容（VR200 約為 GB300 的 2 倍）" if y == 26 else "同上")
for y in (26, 27, 28, 29):
    inp(f"su_{y}", f"非 NVIDIA scale-up 乙太網路交換匣市場 20{y}", "US$B", None, SU[y], None, "Derived", SU_SRC if y == 26 else "同上",
        "Trainium、其他 ASIC、AMD 的 scale-up 交換晶片＋交換匣其他（扣銅纜）；TPU（ICI／OCS）與 NVLink 不計" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"trn_{y}", f"AWS Trainium 顆數 20{y}", "百萬顆", None, NU["trn"][0][j], None, "Derived", f"{rel(NET_F)}｜{NU['trn'][1]}", "網通模型沿用記憶體模型顆數（Assumed：年出貨未找到）" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][0][j], None, "Derived", NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][1], "Recon 用：智邦 AI 營收 ÷ AI IT 資本支出" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、Arista、廣達模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
inp("fx", "匯率（新台幣 ÷ 美元）", "NT$/US$", 29.0, SIL.base("fx_twd"), 33.0, "Assumed", "與 quanta/ 模型相同 32.0", "份額在 FY2026 以同一匯率校準，匯率只在 FY2027 起變動時才影響結果")

sec("B. 乙太網路占比、白牌占比與智邦份額（本專案為權威：白牌交換器廠份額，規則第 13 條）")
inp("eth_26", "乙太網路 ÷ AI 後端交換器（金額）2026", "%", 0.70, 0.75, 0.80, "Assumed", "與 arista/ 模型同值；Dell'Oro 2026-03-10：2025 年乙太網路占 AI 叢集交換器三分之二以上（Interested-party）", "IDC 2Q26：800GbE 占資料中心交換器營收 41.2%（Interested-party）")
inp("eth_27", "乙太網路 ÷ AI 後端交換器 2027", "%", 0.72, 0.80, 0.88, "Assumed", "同上；2027 年無公開預測：未找到", "外推（Assumed）")
inp("eth_28", "乙太網路 ÷ AI 後端交換器 2028", "%", 0.74, 0.83, 0.90, "Assumed", "同上", "—")
inp("eth_29", "乙太網路 ÷ AI 後端交換器 2029", "%", 0.75, 0.85, 0.92, "Assumed", "同上", "—")
inp("wb_26", "白牌（ODM 直供）÷ AI 乙太網路交換器 2026", "%", 0.35, 0.45, 0.55, "Assumed", "Dell'Oro 2025 AI 後端乙太網路：Celestica（白牌）與 NVIDIA 並列第一、合計約 50%，Arista 第三（Interested-party）", "IEEE ComSoc 2026-01：ODM 占資料中心乙太網路交換器約 20–25%（金額）、超大規模資料中心 30–40%（數量；來源未具名）")
inp("wb_27", "白牌 ÷ AI 乙太網路交換器 2027", "%", 0.35, 0.47, 0.57, "Assumed", "超大規模業者自研網路作業系統、擴大白牌（方向）；NVIDIA Spectrum-X 搶占（反方向）", "量化預測：未找到")
inp("wb_28", "白牌 ÷ AI 乙太網路交換器 2028", "%", 0.35, 0.48, 0.58, "Assumed", "同上", "—")
inp("wb_29", "白牌 ÷ AI 乙太網路交換器 2029", "%", 0.35, 0.50, 0.60, "Assumed", "同上", "—")
inp("rel_27", "智邦白牌份額 2027 ÷ 2026（相對值）", "x", 0.80, 1.00, 1.10, "Assumed", "2026 年新增一家超大規模客戶（800G；法人稱約貢獻營收 15%，Interested-party）", "經濟日報 2026-07-09：更多組裝廠加入可能稀釋份額（法人，Interested-party）")
inp("rel_28", "智邦白牌份額 2028 ÷ 2026", "x", 0.70, 0.95, 1.10, "Assumed", "1.6T 世代每代重新競標（Celestica、天弘以外的 ODM 加入）", "—")
inp("rel_29", "智邦白牌份額 2029 ÷ 2026", "x", 0.65, 0.92, 1.10, "Assumed", "同上", "—")

sec("C. AI 加速卡等網路應用產品（AWS Trainium 加速卡模組、L10 等）")
inp("nrel_27", "每顆 Trainium 的智邦營收 2027 ÷ 2026（份額 × 單價，相對值）", "x", 0.65, 0.90, 1.10, "Assumed", "鏡週刊 2026-08-06：智邦自 2Q26 大量出貨 Trainium 3 OAM 模組（Interested-party）", "經濟日報 2026-07-09：更多組裝廠可能稀釋份額；Trainium 4 預計 2H27（法人引述亞馬遜股東信，Interested-party）")
inp("nrel_28", "每顆 Trainium 的智邦營收 2028 ÷ 2026", "x", 0.55, 0.85, 1.10, "Assumed", "同上", "—")
inp("nrel_29", "每顆 Trainium 的智邦營收 2029 ÷ 2026", "x", 0.50, 0.80, 1.10, "Assumed", "同上", "—")
inp("tf", "Trainium 顆數係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 Trainium 顆數為 Assumed（年出貨未找到）", "±20%")

sec("D. 營收結構（FY2025 依公司揭露；FY2026 為 Assumed，用於校準）")
inp("sw25s", "FY2025 網路交換器 ÷ 營收", "%", 0.31, 0.34, 0.36, "Interested-party", FUGLE + "：2025 年前三季網路交換器 34%、網路應用產品 63%、接取 2%（公司法說）", "brodoginvestment 2026-07-11：2025 年資料中心交換器約 31%（作者估計）")
inp("nap25s", "FY2025 網路應用產品（AI 加速卡等）÷ 營收", "%", 0.60, 0.63, 0.65, "Interested-party", "同上：63%（前三季年增 361%）", "brodoginvestment：AI 加速卡與 scale-up 交換器約 63%（作者估計）")
inp("af25", "FY2025 交換器中 AI 資料中心的比例", "%", 0.45, 0.60, 0.75, "Assumed", "公司未揭露交換器的 AI／非 AI 拆分：未找到", "美洲占 2025 年前三季營收 81%（公司）；其餘為企業、電信與前端")
inp("ns26", "FY2026 網路應用產品 ÷ 營收", "%", 0.42, 0.50, 0.58, "Assumed", "美系外資 2026-03（工商時報轉述）：2026 年交換器占營收 48%（Interested-party）→ 網路應用產品約 50%", "工商時報 2026-09-29：1Q26 起高階交換器比重明顯上升、AI 加速卡比重快速上升（Interested-party）")
inp("g_nai26", "非 AI 營收年增 2026", "%", 0.00, 0.10, 0.20, "Assumed", "IDC 2Q26：非資料中心乙太網路交換器年增 15.7%（Interested-party）", "企業、電信、前端交換器與接取；公司未揭露")
inp("g_nai27", "非 AI 營收年增 2027", "%", 0.00, 0.06, 0.12, "Assumed", "同上", "—")
inp("g_nai28", "非 AI 營收年增 2028", "%", 0.00, 0.05, 0.10, "Assumed", "同上", "—")
inp("g_nai29", "非 AI 營收年增 2029", "%", 0.00, 0.05, 0.10, "Assumed", "同上", "—")

sec("E. 毛利率")
inp("d_nap", "網路應用產品毛利率低於交換器與其他的幅度", "pt", 0.03, 0.06, 0.10, "Assumed", "公司：3Q25 毛利率下滑主因產品組合（AI 加速卡占比上升；富果法說摘要）", "工商時報 2026-09-29：AI 加速卡比重上升使毛利率略低於上半年（法人）；產品別毛利率未揭露：未找到")

sec("F. FY2026 錨點（1Q、2Q 實績＋3Q 營收實績＋假設）")
inp("q3_rev", "3Q26 營收（7–9 月月營收合計）", "NT$B", None, 120.74, None, "Verified", CT_Q3 + "：3Q26 營收 1,207.4 億（9 月 403.99 億）", "nstock：1–9 月 2,864.02 億 − 上半年 1,656.58 億＝1,207.44 億")
inp("q3_gm", "3Q26 毛利率", "%", 0.175, 0.185, 0.195, "Assumed", "1Q26 19.55%、2Q26 19.7%（公司）；法人估 3Q26 18.1%（工商時報 2026-09-29，Interested-party）", "時報 2026-10-06：記憶體與 PCB 改代購（客供）使營收認列減少；零組件漲價（Interested-party）")
inp("q4_g", "4Q26 營收季增", "%", -0.05, 0.03, 0.08, "Assumed", "時報 2026-10-06：法人稱 4Q 仍有機會較 3Q 成長（Interested-party）", "公司 4Q 指引：未找到（3Q 法說約 11 月上旬）")
inp("q4_gm", "4Q26 毛利率", "%", 0.17, 0.185, 0.20, "Assumed", "同 3Q 假設", "—")
inp("q_opx", "3Q、4Q26 營業費用 ÷ 營收", "%", 0.037, 0.040, 0.045, "Assumed", "2Q26 4.70%、1Q26 5.22%（公司）；法人估 3Q26 3.9%（Interested-party）", "營收跳增、費用落後")
inp("q_nio", "3Q、4Q26 淨利 ÷ 營業利益", "x", 0.75, 0.80, 0.84, "Derived", "1Q26 8.341 ÷ 10.049＝0.830、2Q26 11.053 ÷ 14.331＝0.771（公司）", "FY2025 0.82（推導）")
inp("q_sh", "股數（十億股；EPS 換算）", "十億股", None, "=19.394/34.7", None, "Derived", UDN_Q2 + "：上半年淨利 193.94 億 ÷ EPS 34.7", "資本額 56.11 億（面額 10 元，鉅亨）＝5.611 億股")

sec("G. 營業費用、業外與稅、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.40, 0.60, 0.80, "Derived", "FY2026 錨點營業費用年增 ÷ 營收年增（Q 頁）", "低／高 Assumed")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.005, 0.0, 0.005, "Assumed", "無", "營益率約 15%，±0.5 個百分點約 ±3% 營業利益")
inp("nio_f", "淨利 ÷ 營業利益", "x", 0.75, 0.80, 0.84, "Derived", "FY2025 0.82、1H26 0.795（公司數字推導）", "匯兌損益與稅率變動")
inp("sh_g27", "股數年變動 2027（現金增資）", "%", 0.0, 0.0178, 0.02, "Interested-party", "經濟日報 2026-10-23（網址見 Sources）：董事會 2026-10-22 通過現增至多 1,000 萬股、稀釋約 1.78%（公司公告轉述）", "發行時程與價格未定：未找到")
inp("sh_g", "股數年變動 2028 起", "%", -0.005, 0.0, 0.01, "Assumed", "股本二十多年大致不變（經濟日報 2026-10-23）", "—")
inp("mf", "AI 交換器市場係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 scale-out 交換器 +90%；智邦訂單能見度到 2028 年（法人轉述）", "±20%")

sec("H. 共識、市場資料（只作 Recon 對帳，不作校準）")
inp("cons_rev26", "共識 FY2026 營收", "NT$B", None, 412.45, None, "Interested-party", FS + "：FactSet 10 位分析師中位數（2026-10-14）", "平均 408.66")
inp("cons_eps26", "共識 FY2026 EPS", "NT$", None, 85.85, None, "Interested-party", "同上（區間 78.59–89.99）", "工商時報 2026-09-29 法人 85.18")
inp("cons_rev27", "共識 FY2027 營收", "NT$B", None, 551.50, None, "Interested-party", "同上", "—")
inp("cons_eps27", "共識 FY2027 EPS", "NT$", None, 120.82, None, "Interested-party", "同上（區間 86.97–139.61）", "工商時報 2026-09-29 法人 107.53")
inp("cons_eps28", "共識 FY2028 EPS", "NT$", None, 170.57, None, "Interested-party", "同上（區間 145.77–205.86）", "—")
inp("br_sw26", "法人：FY2026 交換器 ÷ 營收", "%", None, 0.48, None, "Interested-party", "美系外資 2026-03（旺得富 2026-03-06）", "—")
inp("dell_top2", "2025 AI 後端乙太網路：Celestica＋NVIDIA 合計份額", "%", None, 0.50, None, "Interested-party", "Dell'Oro 2026-03-10", "—")

sec("I. 價格與估值")
inp("px", "股價（2026-10-22 收盤）", "NT$", None, 1775, None, "Interested-party", "鉅亨網 2345（10-22 收盤 1,775、前收 1,815）；臺灣證券交易所 STOCK_DAY（115/10/22 收盤 1,775.00，2026-10-25 查詢）", "永豐金 AI 投資、財報狗同值；工作單：台股無 10-23 收盤時用 10-22（證交所當月資料只到 10/22）；10-23 收盤：未找到")
inp("pe_ttm", "近 12 月 P/E（鉅亨，10-22）", "x", None, 27.88, None, "Interested-party", "鉅亨網", "財報狗 27.89")
inp("tp_wA", "目標價情境權重 A（樂觀）", "%", None, 0.3, None, "Assumed", "Andy 可改", "三個權重合計應為 1")
inp("tp_wB", "目標價情境權重 B（基準）", "%", None, 0.5, None, "Assumed", "Andy 可改", "—")
inp("tp_wC", "目標價情境權重 C（下檔）", "%", None, 0.2, None, "Assumed", "Andy 可改", "—")

sec("J. 目標價參數（v0.2 新增：DCF 與 Comparable 加權、FY2029 後收斂法；黃底＝Andy 可改）")
inp("tp_r", "DCF 折現率（股東權益成本）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25：美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（FY2032 後永續；PL 延伸期 B 情境的收斂目標）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單 2026-10-25）", "Sens 頁 ±1pt（只改 DCF 再投資率與永續成長，不改 PL 收斂路徑）")
inp("tp_gd", "延伸期 A／C 情境長期成長相對 tp_g 的差（A＋、C−）", "pt", None, 0.02, None, "Assumed", "工作單 2026-10-25：A＝tp_g＋2pt、C＝tp_g−2pt", "—")
inp("tp_roe", "近三年平均 ROE（再投資率＝tp_g ÷ tp_roe）", "%", None, 0.4475, None, "Derived", "永豐金 AI 投資 2345 頁（年度 ROE）：2023 39.11%、2024 39.01%、2025 56.13%，平均 44.75%（資料商，Interested-party）", "stockanalysis 近 12 月 ROE 65.58%")
inp("tp_w_dcf", "目標價中 DCF 的權重（其餘為 Comparable）", "%", None, 0.5, None, "Assumed", "Andy 可改（Andy 2026-10-25 11:53 定 50／50）", "Sens 頁 0.3／0.7")
inp("nm10_k", "延伸期 C 情境淨利率目標＝FY2029 淨利率 × 本係數（十年平均未找到）", "x", None, 0.8, None, "Assumed", "工作單：找不到十年平均用最後一年 × 0.8。2016–2021 年淨利率：未找到", "參考：永豐金 AI 投資稅後淨利率 2022 10.58%、2023 10.60%、2024 10.87%、2025 10.59%（四年平均 10.66%）")
inp("cons_tp", "市場共識目標價（FactSet）", "NT$", None, 3500, None, "Interested-party", FS + "：FactSet 10 位分析師目標價 3,500（2026-10-14，鉅亨網）", "stockanalysis（13 位）平均 3,234.92（頁面現價日期不一致，不採）；賣方有目標價誘因，只作對照")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（十億新台幣；EPS 為新台幣；股數為十億股）。口徑：合併營收、稅後淨利（公司公告）。藍字＝輸入；黑字＝公式。實績取自公司季報與月營收公告的媒體轉述（經濟日報、工商時報）"
Q["A2"] = "FY2026 錨點＝1Q26、2Q26 實績＋3Q26 營收實績（月營收）與毛利率、費用假設＋4Q26 假設。3Q26 財報與法說約 11 月上旬（去年 11-07；確切日期未找到），公布後更新 E 欄。FY2025 毛利與營業利益由毛利率 18.09%、營益率 12.94%（永豐、Goodinfo）推導"
hdr(Q, 3, ["項目", "FY2025A", "1Q26", "2Q26", "3Q26（營收實績＋假設）", "4Q26（假設）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：" + GI + "（營收 2,483 億、毛利率 18.1%、營益率 12.9%、EPS 47.13）；1Q26：" + UDN_Q1 + "；2Q26：" + UDN_Q2
S2 = "交叉：永豐金 AI 投資（2025 毛利率 18.09%、營益率 12.94%）；理財周刊 2026-04-21（2025 淨利 263.42 億、EPS 47.13）；工商時報 2026-09-29（2Q26 營業費用 44.91 億）"
def qrow(r, label, vals, fmt="#,##0.00", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
qrow(4, "營收", {"B": 248.32, "C": 70.121, "D": 95.538, "E": "=q3_rev", "F": "=E4*(1+q4_g)"}, tag="Verified／Assumed（4Q）")
qrow(5, "毛利", {"B": "=B4*0.1809", "C": 13.709, "D": 18.823, "E": "=E4*q3_gm", "F": "=F4*q4_gm"}, tag="Verified／Derived（FY2025）／Assumed（3Q、4Q）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEFG"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="1Q26 19.55%、2Q26 19.7%（公司）")
qrow(7, "營業費用", {c: f"={c}5-{c}8" for c in "BCD"} | {"E": "=E4*q_opx", "F": "=F4*q_opx"}, tag="Derived（毛利 − 營業利益）／Assumed")
qrow(8, "營業利益", {"B": "=B4*0.1294", "C": 10.049, "D": 14.331, "E": "=E5-E7", "F": "=F5-F7"}, tag="Verified／Derived")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEFG"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(10, "稅後淨利", {"B": 26.342, "C": 8.341, "D": 11.053, "E": "=E8*q_nio", "F": "=F8*q_nio"}, tag="Verified／Derived（3Q、4Q）")
qrow(11, "淨利 ÷ 營業利益", {c: f"={c}10/{c}8" for c in "BCDEFG"}, fmt="0.000", ann="", tag="Derived", s1="—", s2="1Q26 稅前 103.62 億（業外 +3.1 億；公司）")
qrow(12, "股數（十億股）", {"B": "=26.342/47.13", "C": "=8.341/14.92", "D": "=11.053/19.78", "E": "=q_sh", "F": "=q_sh"}, fmt="0.0000", ann="=q_sh", tag="Derived", s2="淨利 ÷ EPS")
qrow(13, "EPS", {c: f"={c}10/{c}12" for c in "BCDEF"}, ann="=G10/G12", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(14, "  公司揭露 EPS", {"B": 47.13, "C": 14.92, "D": 19.78}, ann="", tag="Verified", s2="永豐金 AI 投資列 2025 年 46.95、1Q26 14.87（口徑不同：疑為追溯調整）")
Q.cell(16, 1, "FY2026 錨點與共識的距離（參考；校準未使用共識）").font = BOLD
Q.cell(17, 1, "FY2026 營收 ÷ 共識 − 1"); Q["B17"] = "=G4/cons_rev26-1"
Q.cell(18, 1, "FY2026 EPS ÷ 共識 − 1"); Q["B18"] = "=G13/cons_eps26-1"
Q.cell(19, 1, "FY2026 營收年增"); Q["B19"] = "=G4/B4-1"
Q.cell(20, 1, "FY2026 營業費用年增"); Q["B20"] = "=G7/B7-1"
Q.cell(21, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B21"] = "=B20/B19"
Q.cell(22, 1, "3Q26 EPS（模型）vs 法人估計 23.91"); Q["B22"] = "=E13/23.91-1"
for rr in range(17, 23): Q[f"B{rr}"].number_format = "0.0%" if rr != 21 else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 10), ("sh", 12), ("eps", 13)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 52
for c in "BCDEFG": Q.column_dimensions[c].width = 15
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 90; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", rel="rel", eth="eth", wb="wb", nrel="nrel", mf="1", tf="1", gn="g_nai", d="d_nap", opxa="0", ns="ns26", af="af25", nio="nio_f")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.0", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    mfx = lambda y, m: "" if y == 2026 else f"*{m}"
    sec_("A. 營收拆分（FY2025 依公司揭露比例；FY2026 校準：網路應用產品比例 Assumed，非 AI 依成長率，AI 交換器為殘差）")
    row("nap0", "網路應用產品（AI 加速卡等）", "NT$B", {"C": f"=a25_rev*nap25s", "D": f"=a26_rev*{P['ns']}"})
    row("nai0", "非 AI（企業、電信、前端交換器、接取）", "NT$B", {"C": f"=a25_rev*(1-nap25s-sw25s*{P['af']})"})
    ws[f"D{R['nai0']}"] = f"=C{R['nai0']}*(1+g_nai26)"; ws[f"D{R['nai0']}"].number_format = "#,##0.0"
    row("aisw0", "AI 資料中心交換器（殘差）", "NT$B", {"C": f"=a25_rev*sw25s*{P['af']}", "D": f"=a26_rev-D{R['nap0']}-D{R['nai0']}"},
        note="Derived：FY2026 交換器合計＝AI＋非 AI 交換器（見 Recon）")
    sec_("B. AI 交換器（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029）")
    row("mkt", "AI 後端 scale-out 交換器市場 × 情境顆數係數", "US$B", {c_(y): f"=mkt_{y % 100}*{umx(y, s)}" + mfx(y, P['mf']) for y in BU}, fmt="#,##0.00")
    row("eth", "乙太網路占比", "%", {c_(y): f"={P['eth']}_{y % 100}" for y in BU}, fmt="0.0%")
    row("su", "非 NVIDIA scale-up 乙太網路交換匣市場 × 係數", "US$B", {c_(y): f"=su_{y % 100}*{umx(y, s)}" + mfx(y, P['mf']) for y in BU}, fmt="#,##0.00")
    row("emk", "AI 乙太網路交換器市場（scale-out × 乙太網路＋scale-up）", "US$B", {c_(y): f"={c_(y)}{R['mkt']}*{c_(y)}{R['eth']}+{c_(y)}{R['su']}" for y in BU}, fmt="#,##0.00")
    row("wb", "白牌占比", "%", {c_(y): f"={P['wb']}_{y % 100}" for y in BU}, fmt="0.0%")
    row("wmk", "白牌 AI 乙太網路交換器市場", "US$B", {c_(y): f"={c_(y)}{R['emk']}*{c_(y)}{R['wb']}" for y in BU}, fmt="#,##0.00")
    row("sh0", "智邦白牌有效份額（FY2026 校準）＝ AI 交換器營收 ÷ 匯率 ÷ 白牌市場", "%", {"D": f"=D{R['aisw0']}/fx/D{R['wmk']}"}, fmt="0.0%",
        note="Derived：分子含前端資料中心交換器，所以高於純後端份額；FY2027 起乘相對值")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "智邦白牌有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("aisw", "AI 交換器營收", "NT$B", {"C": f"=C{R['aisw0']}", **{c_(y): f"={c_(y)}{R['wmk']}*{c_(y)}{R['sh']}*fx" for y in BU}})
    sec_("C. AI 加速卡等網路應用產品")
    row("trn", "AWS Trainium 顆數 × 情境顆數係數", "百萬顆", {c_(y): f"=trn_{y % 100}*{umx(y, s)}" + mfx(y, P['tf']) for y in BU}, fmt="0.00")
    row("vpc0", "每顆 Trainium 的智邦營收（FY2026 校準）", "千新台幣／顆", {"D": f"=D{R['nap0']}/D{R['trn']}"}, fmt="#,##0.0",
        note="Derived：含加速卡模組、L10 等網路應用產品；晶片疑為客供（每顆約 3 千美元）")
    row("nrel", "相對值（份額 × 單價，÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['nrel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("nap", "網路應用產品營收", "NT$B", {"C": f"=C{R['nap0']}", **{c_(y): f"={c_(y)}{R['trn']}*$D${R['vpc0']}*{c_(y)}{R['nrel']}" for y in BU}})
    sec_("D. 非 AI 與合計")
    row("nai", "非 AI 營收", "NT$B", {"C": f"=C{R['nai0']}", "D": f"=D{R['nai0']}", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']}{y % 100})" for y in BU[1:]}})
    row("rev", "營收", "NT$B", {c_(y): f"={c_(y)}{R['aisw']}+{c_(y)}{R['nap']}+{c_(y)}{R['nai']}" for y in FYS})
    row("kns", "非 AI 中的非 AI 交換器比例（固定在 FY2025 結構）", "%", {"C": f"=sw25s*(1-{P['af']})/(1-nap25s-sw25s*{P['af']})"}, fmt="0.0%", note="FY2025：交換器 34% ×（1 − AI 比例）÷ 非 AI 合計；之後各年沿用")
    row("nsw", "  網路交換器：非 AI（企業、電信、前端）", "NT$B", {c_(y): f"={c_(y)}{R['nai']}*$C${R['kns']}" for y in FYS})
    row("sw", "網路交換器營收（AI 資料中心＋非 AI）", "NT$B", {c_(y): f"={c_(y)}{R['aisw']}+{c_(y)}{R['nsw']}" for y in FYS}, note="公司產品別：網路交換器（FY2025 前三季 34%）")
    row("acc", "接取與其他營收（＝非 AI − 非 AI 交換器）", "NT$B", {c_(y): f"={c_(y)}{R['nai']}-{c_(y)}{R['nsw']}" for y in FYS}, note="公司產品別：接取（FY2025 前三季 2%）與其他")
    row("segchk", "檢查：分部合計 − 合併營收（FY2025 實績起）", "NT$B", {"C": f"=C{R['sw']}+C{R['nap']}+C{R['acc']}-a25_rev", **{c_(y): f"={c_(y)}{R['sw']}+{c_(y)}{R['nap']}+{c_(y)}{R['acc']}-{c_(y)}{R['rev']}" for y in FYS[1:]}}, fmt="0.000")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI（交換器＋網路應用產品）÷ 營收", "%", {c_(y): f"=({c_(y)}{R['aisw']}+{c_(y)}{R['nap']})/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("napsh", "網路應用產品 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['nap']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    sec_("E. 損益")
    row("gmo", "交換器與其他毛利率 ＝ 錨點毛利率 ＋ 差距 × 網路應用產品占比（FY2026 校準）", "%", {"D": f"=a26_gp/a26_rev+{P['d']}*D{R['napsh']}"}, fmt="0.00%",
        note="Derived：讓 FY2026 毛利率等於錨點；網路應用產品毛利率＝此值 − 差距")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"=$D${R['gmo']}*({c_(y)}{R['aisw']}+{c_(y)}{R['nai']})+($D${R['gmo']}-{P['d']})*{c_(y)}{R['nap']}" for y in BU[1:]}})
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("opb", "營業費用（基準路徑）", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in BU[1:]}})
    row("opx", "營業費用", "NT$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}})
    row("oi", "營業利益", "NT$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("nio", "淨利 ÷ 營業利益", "x", {"C": "=a25_ni/a25_oi", "D": "=a26_ni/a26_oi", **{c_(y): f"={P['nio']}" for y in BU[1:]}}, fmt="0.000")
    row("ni", "稅後淨利", "NT$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"={c_(y)}{R['oi']}*{c_(y)}{R['nio']}" for y in BU[1:]}})
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", "E": f"=D{r}*(1+sh_g27)", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[2:]}}, fmt="0.0000")
    row("eps", "EPS", "NT$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS}, fmt="#,##0.00")
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 智邦損益（NT$B）。FY2025A 實績、FY2026 錨點（1Q–2Q 實績＋3Q 營收實績＋假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準顆數（顆數係數 1.0；目標價權重視為樂觀）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("aisw", "  AI 交換器營收", "NT$B", "#,##0.0"), ("nap", "  網路應用產品營收", "NT$B", "#,##0.0"),
                                ("nai", "  非 AI 營收", "NT$B", "#,##0.0"), ("aish", "  AI ÷ 營收", "%", "0.0%"), ("gm", "毛利率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"),
                                ("ni", "稅後淨利", "NT$B", "#,##0.0"), ("eps", "EPS", "NT$", "#,##0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收 − 錨點；FY2026 模型毛利 − 錨點（校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：FY2025 營收差（NT$B）"); PL[f"D{sr}"] = f"=C{PB['A']['rev']}-a25_rev"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['gmo']}*(D{PB['A']['aisw']}+D{PB['A']['nai']})+(D{PB['A']['gmo']}-d_nap)*D{PB['A']['nap']}-a26_gp"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 66; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：份額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 智邦在白牌 AI 乙太網路交換器的份額（本專案權威，規則第 13 條）。市場取自 markets/network（scale-out 交換器＋非 NVIDIA scale-up 交換匣）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("emk", "AI 乙太網路交換器市場（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['emk']}", "#,##0.0", "scale-out × 乙太網路占比＋scale-up 交換匣"),
        ("su", "  其中非 NVIDIA scale-up 交換匣", "US$B", lambda y: f"=PL!{FC[y]}{A_['su']}", "#,##0.0", "Trainium、其他 ASIC、AMD"),
        ("wb", "白牌占比", "%", lambda y: f"=PL!{FC[y]}{A_['wb']}", "0.0%", "In 頁 B 節（Assumed）"),
        ("wmk", "白牌 AI 乙太網路交換器市場", "US$B", lambda y: f"=PL!{FC[y]}{A_['wmk']}", "#,##0.0", ""),
        ("sh", "智邦白牌有效份額（含前端資料中心）", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以營收拆分校準"),
        ("aiu", "智邦 AI 交換器營收（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['aisw']}/fx", "#,##0.00", ""),
        ("oth", "其他白牌（Celestica 等）＝ 白牌市場 − 智邦", "US$B", lambda y: f"=PL!{FC[y]}{A_['wmk']}-PL!{FC[y]}{A_['aisw']}/fx", "#,##0.0", "下限：智邦 AI 交換器含前端，所以此列偏低"),
        ("ash", "智邦 ÷ AI 乙太網路交換器市場", "%", lambda y: f"=PL!{FC[y]}{A_['aisw']}/fx/PL!{FC[y]}{A_['emk']}", "0.0%", ""),
        ("trn", "AWS Trainium 顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['trn']}", "0.00", "markets/network Units 頁"),
        ("itr", "智邦 AI 營收 ÷ AI IT 資本支出", "%", lambda y: f"=(PL!{FC[y]}{A_['aisw']}+PL!{FC[y]}{A_['nap']})/fx/itc_{y % 100}", "0.00%", "markets/network 的 AI IT 資本支出")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：FY2026 智邦約占白牌 AI 乙太網路交換器四成上下，與 Dell'Oro「Celestica 為白牌第一」相容的前提是智邦的分子含前端；份額相對值與白牌占比都是 Assumed，是最需要每季檢查的輸入。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」與十年 P/E 規則值（年底落後 P/E，基準為 FY2027 EPS）；目標價為三情境規則值隱含價的加權"
V["A2"] = "循環位置：智邦不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 Arista、廣達；FY(t)＝CY(t)，不平移）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 10  # 第 0 節占用的列數
V.cell(4, 1, "舊法：規則值情境加權（參考）").font = BOLD
hdr(V, 5, ["項目", "值", "相對現價或權重 B", "權重 C", "說明"])
V.cell(6, 1, "規則值情境加權價（參考，非目標價）"); V.cell(7, 1, "相對現價（參考）")
V.cell(8, 1, "情境 A 隱含價（樂觀）"); V.cell(9, 1, "情境 B 隱含價（基準）"); V.cell(10, 1, "情境 C 隱含價（下檔）")
V.cell(11, 1, "情境權重 A／B／C"); V.cell(12, 1, "舊法說明")
V.cell(13, 1, "  權重合計（應為 1）")
for c in range(1, 6): V.cell(6, c).fill = KEY

V.cell(4 + O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5 + O, ["項目", "A", "B", "C", "說明"])
V.cell(6 + O, 1, "現價（2026-10-22 收盤）"); V[f"B{6 + O}"] = "=px"; V[f"B{6 + O}"].number_format = "#,##0"
V.cell(7 + O, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8 + O, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9 + O, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10 + O, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr + O}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr + O}"].number_format = "0.0"
V.cell(11 + O, 1, "未來 12 月 EPS（模型 A）≈ 0.2 × FY2026＋0.8 × FY2027"); V[f"B{11 + O}"] = f"=0.2*PL!D{SUMR[('A', 'eps')]}+0.8*PL!E{SUMR[('A', 'eps')]}"; V[f"B{11 + O}"].number_format = "0.00"
V.cell(12 + O, 1, "  現價 ÷ 未來 12 月 EPS（模型 A）"); V[f"B{12 + O}"] = f"=px/B{11 + O}"; V[f"B{12 + O}"].number_format = "0.0"
V.cell(13 + O, 1, "近 12 月 P/E（鉅亨，10-22）"); V[f"B{13 + O}"] = "=pe_ttm"; V[f"B{13 + O}"].number_format = "0.0"
V[f"E{8 + O}"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"
V.cell(15 + O, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, 16 + O, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for k, s in enumerate(SC):
    i = 17 + O + k
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(21 + O, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 22 + O, ["情境", "FY2027 EPS", "規則基準 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ 基準 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 23 + O + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{17 + O + j}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
# 第 0 節公式
V["B6"] = f"=tp_wA*E{VAL['A']}+tp_wB*E{VAL['B']}+tp_wC*E{VAL['C']}"; V["B6"].number_format = "#,##0"
V["B7"] = "=B6/px-1"; V["B7"].number_format = "0%"
for rr, s in ((8, "A"), (9, "B"), (10, "C")):
    V[f"B{rr}"] = f"=E{VAL[s]}"; V[f"B{rr}"].number_format = "#,##0"
    V[f"C{rr}"] = f"=E{VAL[s]}/px-1"; V[f"C{rr}"].number_format = "0%"
V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0%"
V["B13"] = "=tp_wA+tp_wB+tp_wC"; V["B13"].number_format = "0.00"
V["B12"] = ("P/E：各情境 FY2027 EPS × 十年（2016–2025）年底 P/E 依循環位置取分位（見頂 25 分位、中段中位數、谷底 75 分位；兩個資料商來源分別取分位再平均）；"
            "權重 A 0.3、B 0.5、C 0.2（In 頁，Andy 可改）。偏誤：十年倍數含 2016–2019 年企業網通期的低倍數（約 11–19x），也含 2020、2023–24 年的高倍數期，方向不明；FY2027 三情境 EPS 相同，情境差異只來自循環位置判定")
V["E6"] = "約 FY2027 年底價格、未折現；權重見 In 頁 tp_wA／tp_wB／tp_wC"
V["E8"] = "C 欄＝相對現價"
V.cell(26 + O, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, 27 + O, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = 28 + O
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT = vr
assert vr < 60, vr
PE_CMC = [11.4, 18.6, 15.8, 16.3, 31.7, 28.6, 15.2, 31.7, 35.0, 25.2]       # companiesmarketcap 年底 P/E 2016–2025
PE_GI = [None, None, "=98.6/5.33", "=168/8.91", "=316/9.07", "=260/8.44", "=234.5/14.64", "=523/15.99", "=773/21.49", "=1185/47.13"]  # Goodinfo 年底收盤 ÷ EPS
RB0 = 60
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2016–2025 共 10 年）"},
    {"label": "Goodinfo（年底收盤 ÷ EPS）", "vals": PE_GI, "tag": "Derived（資料商收盤價與 EPS；2018–2025 共 8 年，2016–2017 年底收盤：未找到）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。2016–2019 年為企業網通期（約 11–19x）；2020、2023–2024 年 AI 期重新定價（30–36x）；2025 年 EPS 倍增使年底 P/E 回落到約 25x。")
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額、每顆 Trainium 營收、交換器與其他毛利率在每個案例內依 FY2026 錨點重算，FY2026 營收與毛利恆等於錨點）：對 FY2027、FY2028 的營收、AI 交換器營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 交換器營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 交換器營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("白牌份額流失（低）", "份額相對值 2027 0.80、2028 0.70、2029 0.65（Celestica、新 ODM、NVIDIA 搶占）", dict(rel="RLO")),
         ("白牌份額提升（高）", "份額相對值 1.10", dict(rel="RHI")),
         ("白牌占比 低", "35%（持平）", dict(wb="WBLO")), ("白牌占比 高", "2026 55% → 2029 60%", dict(wb="WBHI")),
         ("乙太網路占比 低", "2026 70% → 2029 75%", dict(eth="ETLO")), ("乙太網路占比 高", "2026 80% → 2029 92%", dict(eth="ETHI")),
         ("AI 交換器市場 ×0.8", "FY2027 起交換器市場 ×0.8（供給受限或內容不如預期）", dict(mf="mf_lo")), ("AI 交換器市場 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("Trainium 份額流失", "相對值 2027 0.65、2028 0.55、2029 0.50（第二供應商、Trainium 4）", dict(nrel="NRLO")),
         ("Trainium 顆數 ×0.8", "FY2027 起 Trainium 顆數 ×0.8", dict(tf="tf_lo")),
         ("網路應用產品占比 42%", "FY2026 拆分：網路應用產品 42%（交換器較大）", dict(ns="ns26_lo")), ("網路應用產品占比 58%", "FY2026 拆分：58%", dict(ns="ns26_hi")),
         ("毛利率差距 3pt", "網路應用產品毛利率只低 3 個百分點", dict(d="d_nap_lo")), ("毛利率差距 10pt", "低 10 個百分點", dict(d="d_nap_hi")),
         ("淨利 ÷ 營業利益 0.75", "匯損或稅率上升", dict(nio="nio_f_lo")),
         ("營業費用率 +0.5pt", "營業費用額外 +0.5% × 營收", dict(opxa="opx_adj_hi"))]
for nm, base, suf in (("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("ETLO", "eth", "_lo"), ("ETHI", "eth", "_hi"), ("WBLO", "wb", "_lo"), ("WBHI", "wb", "_hi"),
                      ("NRLO", "nrel", "_lo")):
    for y in (26, 27, 28, 29):
        if base in ("rel", "nrel", "g_nai") and y == 26: continue
        src = f"{base}{y}{suf}" if base == "g_nai" else f"{base}_{y}{suf}"
        target = f"{nm}{y}" if base == "g_nai" else f"{nm}_{y}"
        wb.defined_names[target] = DefinedName(target, attr_text=wb.defined_names[src].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "aisw", 2027, "#,##0.0"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "aisw", 2028, "#,##0.0"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
assert nn + 3 < 40
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數情境 B、C 的差異從 FY2028 起（交換器與 Trainium 同時縮放）。拆分案例（網路應用產品占比）改變的是 FY2026 校準的起點，FY2026 合計恆等於錨點。")
SENS_NOTE = nn + 3
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識（FactSet 2026-10-14）、法人拆分與研調份額。只作事後對照，不為對上數字而調參數"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "FactSet 中位數", "錨點＝實績＋3Q 營收實績＋4Q 假設")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "FactSet 中位數", "")
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "FactSet 中位數", "")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "FactSet 中位數", "分析師區間 86.97–139.61")
rec("FY2028 EPS vs 共識", "=cons_eps28", pl("eps", 2028), "FactSet 中位數", "")
RC.cell(rr_, 1, "B. 拆分與份額").font = BOLD; rr_ += 1
rec("FY2026 交換器 ÷ 營收（AI＋非 AI 交換器，推導）vs 法人", "=br_sw26",
    lambda s: f"=(PL!D{PB[s]['aisw']}+a25_rev*sw25s*(1-af25)*(1+g_nai26))/PL!D{PB[s]['rev']}", "美系外資 2026-03", "非 AI 交換器假設與非 AI 營收同成長", fmt="0.0%", pt=True)
rec("FY2026 智邦 ÷ AI 乙太網路交換器市場 vs 1 − Celestica、NVIDIA 合計（上限）", "=1-dell_top2", lambda s: f"=PL!D{PB[s]['aisw']}/fx/PL!D{PB[s]['emk']}", "Dell'Oro 2025", "智邦未進前三，應低於上限", fmt="0.0%", pt=True)
rec("FY2027 AI ÷ 營收", None, pl("aish", 2027), "—", "", fmt="0.0%")
rec("FY2027 智邦 AI 營收 ÷ AI IT 資本支出", None, lambda s: f"=(PL!E{PB[s]['aisw']}+PL!E{PB[s]['nap']})/fx/itc_27", "markets/network", "", fmt="0.00%")
rec("FY2026 每顆 Trainium 的智邦營收（千美元）", None, lambda s: f"=PL!D{PB[s]['vpc0']}/fx", "—", "晶片若為客供，約為加速卡＋L10 加工價值", fmt="0.00")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("交換器與其他毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmo']}", "FY2026 錨點", "網路應用產品毛利率＝此值 − d_nap", fmt="0.0%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "", fmt="0.0%", pt=True)
rec("3Q26 EPS（模型）vs 法人估計 23.91", 23.91, lambda s: "=Q!E13", "工商時報 2026-09-29", "", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以智邦 AI 營收 ÷ 部署 GW 計算每 GW 營收"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝季報與月營收公告，本版只經媒體轉述；法說內容經富果摘要）"; SO["A1"].font = BOLD
srcs = [("公司（轉述）：2Q26 財報（營收、毛利、營業利益、淨利、EPS、上半年）", UDN_Q2),
        ("公司（轉述）：1Q26 財報", UDN_Q1),
        ("公司（轉述）：3Q26 月營收（7–9 月）", CT_Q3),
        ("公司（轉述）：3Q25 法說摘要（產品別營收比例、美洲占比）", FUGLE),
        ("公司（轉述）：現金增資 1,000 萬股（2026-10-22 董事會）", UDN_SEO),
        ("二手：工商時報 2026-09-29（法人 3Q26 估計、2026／2027 EPS）", CT_0929),
        ("二手：旺得富 2026-03-06（美系外資：2026 交換器占 48%、新客戶約 15%）", "https://wantrich.chinatimes.com/news/20260306900323-420101"),
        ("二手：經濟日報 2026-07-09（Trainium 3 交換匣、組裝廠稀釋份額）", "https://udn.com/news/story/7253/9617541"),
        ("二手：鏡週刊 2026-08-06（Trainium 3 OAM 模組大量出貨）", "https://www.mirrormedia.mg/external/mirrordaily_77953"),
        ("二手：brodoginvestment 2026-07-11（2025 年資料中心交換器約 31%）", "https://brodoginvestment.substack.com/p/2345"),
        ("共識：FactSet（鉅亨 2026-10-14）", FS),
        ("研調：Dell'Oro 2025 AI 後端（Celestica、NVIDIA 並列第一，Arista 第三）", "https://www.delloro.com/news/ethernet-more-than-doubles-size-of-infiniband-as-the-leading-fabric-for-ai-scale-out-networks-in-2025/"),
        ("研調：IDC 2Q26 乙太網路交換器", "https://www.idc.com/resource-center/blog/ethernet-switch-market-surges-43-4-to-18-9b-in-2q26-as-ai-infrastructure-demand-drives-record-datacenter-spending/"),
        ("二手：IEEE ComSoc 2026-01（ODM 占資料中心交換器比例）", "https://techblog.comsoc.org/2026/01/31/analysis-ethernet-gains-on-infiniband-in-data-center-connectivity-market-white-box-odm-vendors-top-choice-for-ai-hyperscalers/"),
        ("資料商：Goodinfo 歷年經營績效（年底收盤、EPS、毛利率、營益率）", GI),
        ("資料商：永豐金 AI 投資（年度毛利率、營益率、淨利；10-22 收盤）", SINO),
        ("資料商：companiesmarketcap 年底 P/E", "https://companiesmarketcap.com/accton-technology/pe-ratio/"),
        ("市價：鉅亨網 2345（10-22 收盤 1,775、股本 56.11 億）", "https://www.cnyes.com/twstock/2345"),
        ("市場層：markets/network、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：公司原始季報與法說簡報；FY2026 產品別營收比例；交換器的 AI／非 AI 拆分；產品別毛利率；客戶占比（年報 10% 以上客戶）；3Q26 法說日期；10-23 收盤；前瞻 P/E 十年歷史（Bloomberg）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("智邦科技（2345）收支與估值模型 v0.2.1（AI 半導體 Project：網通系統端，白牌交換器 ODM）　2026-10-25", True),
 ("版本：v0.2.1（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）；v0.1（初版）。智邦在白牌 AI 乙太網路交換器的份額在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）。", False),
 ("模型命題：智邦的 AI 交換器營收＝AI 資料中心乙太網路交換器市場（scale-out × 乙太網路占比＋非 NVIDIA scale-up 交換匣；本專案網通模型）× 白牌占比 × 智邦在白牌中的份額；另一半營收是 AWS Trainium 加速卡等網路應用產品（Trainium 顆數 × 每顆營收）；非 AI 另依成長率推。股價約 12 倍模型 FY2027 EPS（共識 FY2027 約 15 倍）：800G → 1.6T 換代與 Trainium 第二供應商之下，份額撐得住嗎？", False),
 ("核心問題：2027–28 年智邦在白牌交換器與 Trainium 加速卡的份額各掉一截，EPS 少多少；現價隱含的路徑合不合理。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029（三年）。金額：十億新台幣；市場以美元、匯率 32.0 換算。口徑：稅後淨利（公司公告）。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 交換器：markets/network 的 scale-out 交換器市場 × 乙太網路占比（Assumed）＋非 NVIDIA scale-up 交換匣（Derived）→ × 情境顆數係數 × 白牌占比（Assumed）× 智邦白牌有效份額（FY2026 校準）× 份額相對值（Assumed）。800G → 1.6T 的單價上升含在網通模型每顆內容；換代的份額風險放在相對值與 Sens。", False),
 ("網路應用產品：Trainium 顆數（網通模型）× 每顆智邦營收（FY2026 校準）× 相對值（Assumed）。非 AI：FY2025 依公司揭露比例推得，FY2026 起依成長率（Assumed）。毛利率：網路應用產品比交換器與其他低 d_nap（Assumed）。", False),
 ("校準：FY2026 ＝ 1Q–2Q 實績＋3Q 營收實績（月營收）＋毛利率、費用與 4Q 假設；FY2026 網路應用產品占比 50%（Assumed）。不用共識（只在 Recon 對帳）。主要客戶 Amazon（Trainium、800G 交換器）與 2026 年新增的另一家美系超大規模業者（法人稱可能為 Meta 等）皆為 Interested-party 說法；客戶占比未找到。", False),
 ("估值：現價 ÷ FY2027、FY2028 EPS；十年 P/E 規則值（兩個資料商來源）× FY2027 EPS；目標價＝三情境規則值隱含價加權（0.3／0.5／0.2）。循環位置以本模型營收最高年判定。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文只經媒體轉述（季報、月營收、法說摘要）。", False),
 ("", False),
 ("已知限制：(1) FY2026 產品別拆分（網路應用產品 50%）與交換器的 AI 比例為 Assumed，決定 AI 交換器與 Trainium 兩條線的起點；(2) 白牌占比與份額相對值皆為 Assumed；(3) 網通模型 Trainium 顆數為 Assumed（年出貨未找到）；(4) 記憶體與 PCB 改客供（代購）會壓低營收、墊高毛利率，本版未拆；(5) 3Q26 財報未出，3Q 毛利率與費用為假設；(6) 十年 P/E 兩來源 2016–2017 年只有一個；(7) R4 產業部署路徑未建（Recon 預留列）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ---------------- v0.2：三情境延伸到 FY2032、Val 第 0 節（DCF 與 Comparable 加權）、Sens 目標價參數 ----------------
XH, XR, XC = TP.extend_pl(PL, PB, FC, FYS, 2029, {"rev": "rev", "ni": "ni", "sh": "sh_", "eps": "eps"}, "NT$",
    "三情境延伸到 FY2032（v0.2；FY2025A–FY2029 引用上方推導區塊，FY2030 起為收斂法，非驅動因子推導；收斂年度為橘底）",
    "上游市場層（markets/network、台積電顆數）只到 CY2029＝FY2029；FY2030 起：營收年增由 FY2029 年增線性收斂到 FY2032 的長期成長（A＝tp_g＋tp_gd、B＝tp_g、C＝tp_g−tp_gd），"
    "淨利率由 FY2029 線性收斂到中段利潤率（A＝FY2029、B＝FY2029 與 C 目標的中點、C＝FY2029 × nm10_k；十年平均未找到），股數沿用 FY2029 年變動")
_h = next(r for r in range(1, V.max_row + 1) if V.cell(r, 1).value == "P/E：十年歷史倍數")
TPL, TPN = TP.add_tp_section(wb, V,
    eps={s: [f"PL!{XC[y]}{XR[(s, 'eps')]}" for y in range(2027, 2033)] for s in SC}, e_mon=12,
    fwd={s: [(f"PL!{XC[2027]}{XR[(s, 'eps')]}", "2/12"), (f"PL!{XC[2028]}{XR[(s, 'eps')]}", "10/12")] for s in SC},
    fwd_note="會計年度＝日曆年：約 12 個月後（2027-10 底）的未來 12 個月＝2027-11～2028-10（2/12 × FY2027＋10/12 × FY2028）",
    peers=[("Arista Networks", "ANET", 37.06, "2026-10-23", "stockanalysis.com 統計頁（Forward PE）", "Interested-party（資料商彙整賣方預估）"),
           ("Celestica", "CLS", 16.65, "2026-10-23", "同上", "Interested-party（同上）"),
           ("緯穎", "6669.TW", 11.73, "2026-10-21", "同上（台股頁，價格延遲）", "Interested-party（同上）"),
           ("廣達", "2382.TW", 11.52, "2026-10-21", "同上（台股頁，價格延遲）", "Interested-party（同上）")],
    peer_note="法人常用的資料中心網通與 AI 伺服器 ODM 同業（Arista 為品牌交換器、Celestica 為白牌交換器直接競爭者、緯穎與廣達為台灣 AI 伺服器 ODM）；不含 A 股（Andy 11:53）。前瞻 P/E 為資料商的未來 12 個月共識 EPS",
    cons_tp="cons_tp", cons_note="FactSet：10 位分析師目標價（2026-10-14，鉅亨網轉述）；Interested-party",
    cons_fwd="2/12*cons_eps27+10/12*cons_eps28", cons_fwd_note="同權重的共識 EPS：2/12 × FactSet FY2027 中位數 120.82＋10/12 × FY2028 中位數 170.57",
    eps28={s: f"PL!{XC[2028]}{XR[(s, 'eps')]}" for s in SC}, hist=f"B{_h + 2}:B{_h + 11}", fmt="#,##0",
    method="DCF：股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），tp_r 10%、tp_g 4%、tp_roe 近三年平均 44.75%（再投資率約 9%）；基準日 2026 年底，FY2027 起逐年折現，FY2032 後永續。"
           "Comparable：Arista、Celestica、緯穎、廣達前瞻 P/E 中位數 × 12 個月後前瞻 EPS。偏誤方向：伺服器 ODM（緯穎、廣達）毛利率與成長低於智邦，中位數偏低（Comparable 偏保守）；"
           "C 情境 FY2030 起淨利率收斂到 FY2029 × 0.8（Assumed）；歷史分位用十年年底落後 P/E（companiesmarketcap）；永續價值占 DCF 價值過半（Sens 頁）。")
assert nn + 9 < 40, nn
TP.add_tp_sens(SE, nn + 5, TPL, "#,##0")
TP.dump_rowmap(TPL, TPN, XH, XR, XC, nn + 5)
TP.yellow_inputs(IN)
RD.cell(len(lines) + 1, 1, "v0.2（2026-10-25）：(1) 目標價改為 DCF 與 Comparable 加權（Val 頁最前面「0. 目標價」：0.1 DCF 股東權益現金流、0.2 Comparable 同業前瞻 P/E、0.3 結論；Val 既有內容整體下移，公式一併改寫）；"
        "v0.1 的「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留。(2) PL 頁末新增三情境延伸到 FY2032：FY2030 起為收斂法，非驅動因子推導（上游市場層只到 CY2029＝FY2029）。"
        "(3) 現價維持 2026-10-22 收盤 1,775（10-23 收盤未取得；補證交所為第二來源）。(4) In 頁 J 節新增目標價參數（黃底，Andy 可改）；Sens 頁新增 tp_r、tp_g、tp_w_dcf 對加權目標價的影響。FY2025–FY2029 既有數字不變。")
RD.cell(len(lines) + 2, 1, "情境對應（v0.2 延伸期）：A＝順利（基準顆數）、B＝平平（CY2028 ×0.8、CY2029 ×0.85）、C＝不如預期（CY2028–29 ×0.6）；FY2027 三情境相同，情境差從 FY2028 起。本版未做先驗，無 Andy 合理 P/E 對照。")

# ---------------- v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 14:00、14:17；既有數值與儲存格位置不變） ----------------
_UPW = [r_ for r_ in range(1, 60) if V.cell(r_, 1).value == "上檔空間：加權目標價"]
assert _UPW == [42], _UPW
GR_ROW, RC_ROW = 51, 52   # 0.3 結論後的空白預留列（不移動其他列）
for _r in (GR_ROW, RC_ROW):
    assert all(V.cell(_r, _c).value is None for _c in range(1, V.max_column + 1)), _r
V.cell(GR_ROW, 1, "上檔空間級距")
V.cell(GR_ROW, 2, '=IF(B42>0.5,"大於 50%",IF(B42>0.25,"25%–50%",IF(B42>=0,"0%–25%",IF(B42>=-0.25,"0%～−25%",IF(B42>=-0.5,"−25%～−50%","小於 −50%")))))').number_format = "@"
V.cell(GR_ROW, 5, "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
V.cell(RC_ROW, 1, "投資建議")
V.cell(RC_ROW, 2, '=IF(B42>0.5,"強力買進",IF(B42>0.25,"買進",IF(B42>=-0.25,"中立",IF(B42>=-0.5,"放空","強力放空"))))').number_format = "@"
V.cell(RC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = [
       "目標價 NT$2,275（DCF 2,584、Comparable 1,966 各半），比現價 1,775（2026-10-22）高 28.2%，級距 25%–50%；Comparable 單獨為 +10.8%。主因是 DCF 不是 EPS：本模型 12 個月後前瞻 EPS 138.6 比共識前瞻 EPS 162.3 低 14.6%，在同業前瞻 P/E 中位數 14.2 倍下 Comparable 只高一成；DCF 單獨 +45.6%，因 ROE 44.8% 遠高於折現率 10%（可發放比例 91%），而現價只有 FY2027 EPS 的 12.0 倍。共識目標價 3,500 隱含 21.6 倍，比同業高 52.0%。",
       "FY2028 EPS 三情境分歧大：A 169.2（與 FactSet 中位數 170.57 相近）、B 133.8（−20.9%）、C 95.7（−43.4%），差距來自 CY2028 起的顆數情境；單一變數中白牌份額流失（相對值 0.80／0.70）FY2028 EPS −18.2%、AI 交換器市場 ×0.8 −13.8%、Trainium 份額流失 −10.4%。",
       "追蹤 3Q26 財報與法說（去年在 11-07，約 2026-11 上旬）與 10 月營收（11-10 前）：本版 3Q 毛利率與費用為假設，FY2026 網路應用產品占比 50% 為 Assumed（42%／58% 時 FY2028 EPS +6.3%／−6.7%）。網通 CPO 目前只有一條路徑、三情境相同，1.6T 換代的 CPO 份額風險待 markets/network 改版。"
      ]   # 重點觀察（與 README.md 相同文字）
_n0 = RD.max_row
RD.cell(_n0 + 1, 1, "v0.2.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（Val 頁 0.3 結論加「上檔空間級距」「投資建議」兩列、README 頁加重點觀察三點；數值不變，既有儲存格位置不變）。")
RD.cell(_n0 + 2, 1, f'="投資建議："&Val!B{RC_ROW}&"（依上檔空間固定規則）"').font = BOLD
for _i, _t in enumerate(OBS, 1): RD.cell(_n0 + 2 + _i, 1, f"重點觀察 {_i}：" + _t)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# 相對倍數調整（Andy 2026-10-27「照建議做」；tools/relative_multiple.py，2026-10-29 接入）
import os, sys  # noqa: E401
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools"))
import relative_multiple as RMX  # noqa: E402
RMX.apply(wb, "accton")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tools", "in_layout")); import apply_in_layout  # noqa: E402
apply_in_layout.apply(wb, "v0.3.1", "2026-10-30")  # v0.3.1：In 頁版面（個股模型原則第 8 節；數值不變）
wb.save(sys.argv[1])
if os.environ.get("ACCTON_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "UN_NOTE": UN_NOTE, "CHK": CHK_ROW, "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE,
               "VAL": VAL, "WMT": WMT, "RC_LAST": RC_LAST, "block_end": block_end, "rule_rows": rule_rows},
              open(os.environ["ACCTON_ROWMAP"], "w"), ensure_ascii=False, indent=1)
