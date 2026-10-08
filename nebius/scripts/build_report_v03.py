# v0.2a 步驟 8：對照 Excel docs/reports/20261008_nebius_v0.3_收入錨定.xlsx（數值取自 tkanchor.py、tk_sens.json、
# out/attrib_tkanchor.json（node scripts/attrib_tkanchor.js）與 dist/ v0.3 成品；只作報告，不進模型）。
# 用法：python3 scripts/build_report_v03.py [輸出.xlsx]
import json, os, sys, subprocess
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, REPO)
import calendar_q, tkanchor
co = calendar_q.load(REPO); _, tk = tkanchor.load(co); P = co['periods']; K3 = ['low', 'base', 'high']
LB = co['scenarios']['labels']; AM = co['pricing']['anchorMultiple']
TS = json.load(open(os.path.join(REPO, 'tk_sens.json'), encoding='utf-8'))
AT = json.load(open(os.path.join(REPO, 'out', 'attrib_tkanchor.json'), encoding='utf-8'))
dist = [f for f in os.listdir(os.path.join(REPO, 'dist')) if f.endswith('.xlsx')][0]
X = load_workbook(os.path.join(REPO, 'dist', dist), data_only=True)
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864'); B = Font(bold=True)
wb = Workbook(); first = True
def sheet(name, head, rows, note=None):
    global first
    ws = wb.active if first else wb.create_sheet(); first = False; ws.title = name
    r = 1
    if note: ws.cell(1, 1, note).font = Font(italic=True, size=9); r = 2
    for j, h in enumerate(head): c = ws.cell(r, 1 + j, h); c.font = H; c.fill = HF
    for row in rows:
        r += 1
        for j, v in enumerate(row): ws.cell(r, 1 + j, v)
    ws.column_dimensions['A'].width = 48
    for col in 'BCDEFGHIJ': ws.column_dimensions[col].width = 14
R = {sc: tkanchor.compute(co, tk, sc, co['periodYears']) for sc in K3}
rows = []
for sc in K3:
    rows.append([f'{LB[sc]}｜v0.2（舊三情境）'] + [x * 1e3 for x in co['scenarios']['revMW'][sc]])
    rows.append([f'{LB[sc]}｜v0.3（錨 × k，價格基準）'] + [x * 1e3 for x in R[sc]['rev']])
    rows.append([f'{LB[sc]}｜變動'] + [R[sc]['rev'][i] * 1e3 - co['scenarios']['revMW'][sc][i] * 1e3 for i in range(5)])
sheet('每MW收入_v0.2對v0.3', ['US$m／MW-IT·年'] + P, rows, '每 MW 年收入（100% 計費時數）；v0.2＝路徑 A 成本加成與路徑 B 市場價 50/50 平均、隨容量情境同向（已停用）')
rows = []
for sc in K3:
    for key, lab, f in (('anchor', '錨（IF_HoldEcon 在役世代加權，US$m）', 1e3), ('ls', '長約占比', 1), ('k', '定價倍數 k', 1), ('rev', '錨 × k（US$m）', 1e3), ('capRatio', '上限比（÷ IF_RevGWFleet）', 1)):
        rows.append([f'{LB[sc]}｜{lab}'] + [x * f for x in R[sc][key]])
    for g in co['fleet']['generations']: rows.append([f'{LB[sc]}｜平均在役占比｜{g}'] + R[sc]['share'][g])
sheet('錨與k逐年', ['項目'] + P, rows, f"Tokenomics {co['tokenomics']['version']}（{co['tokenomics']['currentFile']}，合併 {co['tokenomics']['mergeCommit']}）；k_長約 {AM['long']['base']}、k_現貨 {AM['spot']['base']}")
rows = [[e['label'], e['gen'], e['price'], e['unit'], e['tkName'], tk[e['tkName']]['values'][e['gen']]['基準'], e['price'] / tk[e['tkName']]['values'][e['gen']]['基準'], e['use'], e['contract'], e['term'], e['tag'], e['source'], e['date'], e['url'], e['note']] for e in AM['evidence']]
rows += [[f"合約組合｜{x['label']}", '', x['value'], x['unit'], '', '', '', '', '', '', x['tag'], x['source'], x.get('date', ''), x['url'], ''] for x in AM['contractMix']]
rows += [[f"長約合約｜{c['label']}", '', c['mw'], 'MW', '', '', '', f"起始期 {c['start']}", f"下緣 {c['lo']}", f"上緣 {c['hi']}", c['tag'], c['source'], '', c['url'], c['basis']] for c in AM['longShare']['contracts']]
rows += [['找不到', '', '', '', '', '', '', '', '', '', '', t, '', '', ''] for t in AM['notFound']]
sheet('k證據表', ['證據', '世代', '價格', '單位', 'Tokenomics 名稱', 'Tokenomics 基準', '倍數 k', '用途', '合約型態', '期間', '標記', '來源', '日期', '網址', '說明'], rows)
ws = X['每MW收入_錨定']; q2 = []
on = False
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if a and str(a).startswith('F｜'): on = True; continue
    if a and str(a).startswith('G｜'): break
    if on and a and a not in ('項目',): q2.append([a, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 9).value])
sheet('Q2驗證', ['項目', '單位', '數值', '說明'], q2, '只作驗證、不校準 k（已決定事項 14）；(i)＋(ii)＋(iii)＋(iv)＝總差距')
rows = [[LB[sc]] + [TS['matrix'][sc][px]['tgt'] for px in K3] + [TS['matrix'][sc][px]['eq'] for px in K3] + [TS['matrix'][sc][px]['gap'] for px in K3] for sc in K3]
sheet('容量x價格3x3', ['容量情境', '目標價｜價格低', '目標價｜價格基準', '目標價｜價格高', '股權募資｜低', '股權募資｜基準', '股權募資｜高', '融資前缺口｜低', '融資前缺口｜基準', '融資前缺口｜高'], rows,
      f"價格低／基準／高＝k_長約 {AM['long']['low']}／{AM['long']['base']}／{AM['long']['high']}、k_現貨 {AM['spot']['low']}／{AM['spot']['base']}／{AM['spot']['high']}；US$／股、US$bn")
rows = [['基準（目前輸入）'] + [TS['matrix'][sc]['base']['tgt'] for sc in K3] + [TS['matrix'][sc]['base']['eq'] for sc in K3] + [TS['matrix'][sc]['base']['gap'] for sc in K3]]
rows += [[x['label']] + [x['res'][sc]['tgt'] for sc in K3] + [x['res'][sc]['eq'] for sc in K3] + [x['res'][sc]['gap'] for sc in K3] for x in TS['sens']]
sheet('敏感度', ['變動', '目標價｜保守', '目標價｜基準', '目標價｜積極', '股權募資｜保守', '股權募資｜基準', '股權募資｜積極', '融資前缺口｜保守', '融資前缺口｜基準', '融資前缺口｜積極'], rows)
rows = []
for sc in K3:
    a = AT[sc]; S = a['steps']
    rows.append([LB[sc], S[0]['tgt'], 0, 0, 0, a['d1'], a['d2'], a['d3'], 0, S[3]['tgt'], a['total'], S[0]['gap'], S[3]['gap'], S[0]['eq'], S[3]['eq']])
sheet('目標價變動拆解', ['情境', 'v0.2', '(a) 時間推移', '(b) 實際數', '(c) 假設', '(d)① 錨取代舊推導（k＝1）', '(d)② 套用 k', '(d)③ 情境軸分離', '(d)④ 上限檢查', 'v0.3', '總變動', '融資前缺口 v0.2', '融資前缺口 v0.3', '股權募資 v0.2', '股權募資 v0.3'], rows,
      '已決定事項 12；(a)(b)(c) 由 scripts/attrib.py（Excel）確認為 0；(d) 再拆由 scripts/attrib_tkanchor.js（HTML 引擎，與 Excel 經 cmp31 一致）')
G = co['fleet']['generations']; rows = []
for n in ('IF_OpexGW', 'IF_PowerCost', 'IF_MaintIT', 'IF_MaintFac', 'IF_StaffSW', 'IF_TaxIns'):
    rows.append([f'Tokenomics {n}（在役世代加權，基準情境，US$m/MW）'] + [sum(R['base']['share'][g][t] * tk[n]['values'][g]['基準'] for g in G) for t in range(5)])
op = rows[0][1:]
rows.append(['每 MW 年收入（錨 × k，基準）'] + [x * 1e3 for x in R['base']['rev']])
rows.append(['由下而上隱含 EBITDA 率（未扣管銷、租金）'] + [1 - op[t] / (R['base']['rev'][t] * 1e3) for t in range(5)])
rows.append(['模型 EBITDA 率路徑（現行，可觀察 neocloud 區間）'] + [co['defaults']['ebStart'] + (co['defaults']['ebSteady'] - co['defaults']['ebStart']) * t / 4 for t in range(5)])
sheet('成本端評估', ['項目'] + P, rows, '只評估並列、不切換（工作單預設 7）；是否切換由 Andy 決定。穩態 EBITDA 率 59%／70%／75% 時基準目標價 $132.2／$203.7／$236.0')
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, 'docs', 'reports', '20261008_nebius_v0.3_收入錨定.xlsx')
wb.save(out); print('saved', out)
