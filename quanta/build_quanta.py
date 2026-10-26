# 廣達電腦（2382）收支與估值模型 v0.2（AI 半導體 Project：伺服器組裝次產業樣板）
# 用法：python3 quanta/build_quanta.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；整櫃價格與市場以美元輸入、以匯率 fx 換算；EPS：新台幣。
# 口徑：歸屬母公司淨利。業外收益與稅前利益：公開摘要未找到，以「淨利 ÷ 營業利益」一個係數涵蓋業外、稅與非控制權益。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ITC = None
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    if r[0].value == "AI IT 資本支出（加速器 × 每顆 IT capex）":
        ITC = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert ITC and all(v is not None for v in ITC[0]), ITC
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("racks", "capit", "rack"):
        BR[k] = (r[4].value, r[7].value, f"{rel(BR_F)}｜In!E{r[0].row}（GB300）、H{r[0].row}（VR200）")
assert len(BR) == 3, BR.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B），整櫃價格與市場為美元；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝廣達法說（經富果、優分析、Taipei Times 摘要）、月營收公告；公司的 AI 占比與展望為 Interested-party；券商估計皆為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
FUGLE = "https://blog.fugle.tw/post/earnings-call-2382-2026-08-13"
UAN = "https://uanalyze.com.tw/articles/5252953366"
KGI = "https://www.kgi.com.hk/zh-cn/-/media/files/kgishk/research-reports/tw-reports/quanta-computer-2382-tt-20260813en.pdf"
SD1008 = "https://statementdog.com/news/17848"
CTEE = "https://www.ctee.com.tw/news/20261008700747-430201"
UDN0910 = "https://money.udn.com/money/amp/story/5607/9745501"
UAN25 = "https://uanalyze.com.tw/articles/3341848159"

sec("A. AI 部署規模（讀 markets/network、markets/chip_bridge、tsmc 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, ITC[0][j], None, "Derived", ITC[1],
        "加速器顆數 × 每顆 IT capex（含 ASIC）；本模型以此除以每 GW IT 資本支出得到部署 GW（與台達電模型同口徑）" if y == 26 else "同上")
inp("itpg_gb", "每 GW IT 資本支出（GB300）", "US$B", None, BR["capit"][0], None, "Derived", BR["capit"][2], "Tokenomics v20.1 IF_CapexIT（經橋接表；唯讀）")
inp("itpg_vr", "每 GW IT 資本支出（VR200；2027 起）", "US$B", None, BR["capit"][1], None, "Derived", BR["capit"][2], "同上；Rubin Ultra 與 ASIC 機架未另訂（Assumed：同 VR200）")
inp("vr26", "2026 年部署中 Rubin 世代占比（混合用）", "%", 0.15, 0.22, 0.30, "Interested-party", "TrendForce 2026 組合 Rubin 22%（經 markets/memory）", "廣達：Vera Rubin 4Q26 小量、1Q27 放量（財報狗 2026-10-08）")
inp("racks_gb", "每 GW 機架數（GB300）", "架", None, BR["racks"][0], None, "Derived", BR["racks"][2], "Tokenomics IF_RacksPerGW（經橋接表）")
inp("racks_vr", "每 GW 機架數（VR200）", "架", None, BR["racks"][1], None, "Derived", BR["racks"][2], "同上")
inp("rack_gb", "整櫃買方價格（GB300，Tokenomics）", "US$M/架", None, BR["rack"][0], None, "Derived", BR["rack"][2], "Recon 用：ODM 價 ÷ 買方價")
inp("rack_vr", "整櫃買方價格（VR200，Tokenomics）", "US$M/架", None, BR["rack"][1], None, "Derived", BR["rack"][2], "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、Arista、台達電模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
inp("fx", "匯率（新台幣 ÷ 美元）", "NT$/US$", 30.5, 32.0, 33.0, "Derived", "與台達電模型相同（公司換算隱含約 32.7）", "先驗比對用 32.7（只影響美元換算的市場與每 GW）")

sec("B. 組裝價格、廣達份額與客供（本專案為權威：NVIDIA 整櫃組裝份額、每櫃 ODM 價格、客供比例，規則第 13 條）")
inp("odm_gb", "ODM 每櫃價格 GB300", "US$M/架", 3.6, 4.0, 4.3, "Interested-party", CTEE + "：GB300 ASP 約 US$4M（法人，未具名）", "FTNN 2026-10-09 同值；上限＝Tokenomics 買方價 4.3")
inp("odm_vr", "ODM 每櫃價格 VR200", "US$M/架", 6.5, 6.75, 7.0, "Interested-party", CTEE + "：VR ASP 估 US$6.5–7M（法人）", "FTNN 2026-10-09 同值；Tokenomics 買方價 8.4")
inp("vpg_28", "每 GW ODM 整櫃價值 2028 ÷ 2027（Rubin Ultra）", "x", 0.85, 1.0, 1.2, "Assumed", "Rubin Ultra（Kyber）每櫃價格與每 GW 架數：未找到", "每 GW 價值 GB300 約 27.1B → VR200 約 27.4B（4.0 × 6,764 vs 6.75 × 4,052）：換世代時每 GW 大致持平")
inp("vpg_29", "每 GW ODM 整櫃價值 2029 ÷ 2027", "x", 0.85, 1.0, 1.2, "Assumed", "同上", "—")
inp("rel_27", "廣達 AI 伺服器有效份額 2027 ÷ 2026（相對值）", "x", 0.85, 1.0, 1.1, "Assumed", UDN0910 + "：美系券商 8 月調查 鴻海約 2,600 櫃、廣達 2,500–2,550、緯創 2,300–2,400（Interested-party，單一來源）", "公司：前五大 CSP 為主，Neocloud GPU 與 CSP ASIC 專案小量出貨（2Q26 法說）；Neocloud 繞過 OEM 找 ODM（FTNN 2026-10-09）")
inp("rel_28", "廣達份額 2028 ÷ 2026", "x", 0.75, 0.95, 1.1, "Assumed", "同上；台灣組裝產能 2028 再翻倍（公司）", "—")
inp("rel_29", "廣達份額 2029 ÷ 2026", "x", 0.70, 0.90, 1.1, "Assumed", "同上", "—")
inp("cs_27", "客供比例 2027（AI 毛額價值中改為客戶供料、不計營收的比例）", "%", 0.0, 0.10, 0.25, "Assumed", FUGLE + "：3Q26 起部分高價關鍵料件改客供，GPU 與 ASIC 客戶都有；比重仍低（公司）", KGI + "：2H26 客供比重小，營益率持平 1H26")
inp("cs_28", "客供比例 2028", "%", 0.0, 0.20, 0.40, "Assumed", "同上；公司與客戶協商轉換以提高利潤率（Taipei Times 2026-05-15）", "—")
inp("cs_29", "客供比例 2029", "%", 0.0, 0.25, 0.50, "Assumed", "同上", "—")
inp("keep", "客供部分保留的毛利比例（相對買賣模式）", "x", 0.70, 0.90, 1.0, "Assumed", "公司：客供可減輕毛利率稀釋與營運資金壓力（2Q26 法說）", "買賣模式下轉手 GPU 的加價很薄；客供後少掉這部分加價：未找到量化")

sec("C. 營收結構錨點（公司揭露比例 × 營收）")
inp("srv_26", "2026 年伺服器營收 ÷ 營收", "%", 0.82, 0.85, 0.88, "Interested-party", FUGLE + "：伺服器自 1Q26 起超過 80%（公司）", KGI + "：雲端伺服器 2026F 85%、2027F 88%")
inp("ais_26", "2026 年 AI 伺服器 ÷ 伺服器營收", "%", 0.75, 0.80, 0.82, "Interested-party", UAN + "：1H26 75–80%，全年約 80%（公司）", KGI + "：同")
inp("g_nai26", "2026 年非 AI 營收年增（反推 FY2025 AI 營收）", "%", -0.05, 0.05, 0.15, "Assumed", "公司：筆電雙位數衰退、一般伺服器雙位數成長（2Q26 法說、財報狗 10-08）", SD1008 + "：公司維持 AI 伺服器「三位數年增」；若只倍增，非 AI 須 +67%，與筆電衰退不符，所以倍增讀為下限")
inp("g_nai27", "非 AI 營收年增 2027", "%", -0.05, 0.03, 0.08, "Assumed", "筆電 2026 年雙位數衰退、低於 20%（公司，財報狗 10-08）；一般伺服器雙位數成長（公司）", "筆電約占 2026 營收 11–13%")
inp("g_nai28", "非 AI 營收年增 2028", "%", -0.05, 0.02, 0.06, "Assumed", "同上", "—")
inp("g_nai29", "非 AI 營收年增 2029", "%", -0.05, 0.02, 0.06, "Assumed", "同上", "—")

sec("D. 毛利率")
inp("d_ai", "AI 伺服器毛利率（毛額基礎）減其他營收毛利率", "pt", -0.035, -0.02, -0.005, "Assumed", "公司：高價 AI 機櫃稀釋毛利率；FY2025 6.98% → 1H26 4.92%（AI 占比上升）", "部門別毛利率未揭露：未找到")
inp("vrd", "VR 世代起 AI 毛利率（毛額基礎）變動（FY2027 起）", "pt", -0.006, -0.003, 0.0, "Assumed", "凱基 2026-08-13（網址見 Sources）：VR 每櫃價格較高，毛利金額增加、毛利率可能被稀釋；凱基毛利率 2026F 4.9% → 2027F 4.7%", "每櫃加工費與每櫃價格不同比例上升：未找到量化")

sec("E. FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設）")
inp("q3_rev", "3Q26 營收（月營收合計）", "NT$B", None, 1115.5, None, "Verified", SD1008 + "：3Q26 1.1155 兆、前三季 2.9614 兆", "TechNews 2026-10-08：前 9 月 2.96 兆")
inp("q4g", "4Q26 營收季增", "%", -0.15, -0.07, 0.05, "Interested-party", SD1008 + "：法人估全年接近 4 兆（4Q 約 1.04 兆，季減約 7%）", "公司：GB300 為 4Q 主要動能、VR 4Q 小量；9 月月減 23%（出貨時點）")
inp("q_gm", "3Q、4Q26 毛利率", "%", 0.047, 0.050, 0.053, "Interested-party", FUGLE + "：全年利潤率維持 1H26 水準；2Q26 5.02%（公司）", KGI + "：2026F 4.9%")
inp("q_opx", "3Q、4Q26 營業費用 ÷ 營收", "%", 0.0165, 0.0175, 0.019, "Assumed", "2Q26 1.72%（新低）、1H26 1.81%（公司）", "—")
inp("q_nio", "3Q、4Q26 淨利 ÷ 營業利益", "%", 0.84, 0.87, 0.89, "Derived", "1H26 49.84 ÷ 57.3＝0.870；FY2025 74.98 ÷ 87.5＝0.857（公司）", "業外收益與稅前利益：未找到（公開摘要未列）")
inp("q_sh", "股數（十億股；EPS 換算）", "十億股", None, "=49.84/12.93", None, "Derived", "1H26 淨利 49.84 ÷ EPS 12.93（公司）", "FY2025 74.98 ÷ 19.45＝3.855")

sec("F. 營業費用、淨利係數、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.2, 0.4, 0.7, "Derived", "FY2026 錨點實測（Q 頁 B26）", "營業費用率 FY2025 2.9% → 2Q26 1.72%")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.002, 0.0, 0.002, "Assumed", "無", "營益率約 3%，±0.2 個百分點已是約 ±7% 的營業利益")
inp("nio_f", "淨利 ÷ 營業利益", "%", 0.75, 0.865, 0.89, "Derived", "同 q_nio（1H26 0.870、FY2025 0.857）", "低值＝券商隱含：凱基 FY2026 EPS 27.06 × 3.855 ÷（3.2% × 4,219）≈ 0.77、FY2027 ≈ 0.73（原因未找到）")
inp("sh_g", "股數年變動", "%", 0.0, 0.0, 0.005, "Assumed", "FY2025 3.855 → 1H26 3.855（公司 EPS 與淨利）", "—")
inp("mf", "AI 部署係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 AI IT 資本支出約 2 倍", "±20%")

sec("G. 共識、市場資料（只作 Recon 對帳，不作校準）")
inp("b_rev26", "券商 FY2026 營收（凱基 08-13，3Q 前）", "NT$B", None, 4219.196, None, "Interested-party", KGI, "3Q 只達法人預期 92%（工商時報 10-08）")
inp("b_eps26", "券商 FY2026 EPS（凱基 08-13）", "NT$", None, 27.06, None, "Interested-party", KGI, "—")
inp("b_rev27a", "券商 FY2027 營收（凱基 08-13）", "NT$B", None, 5220.407, None, "Interested-party", KGI, "—")
inp("b_rev27b", "券商 FY2027 營收（法人，10-08）", "NT$B", None, 5300.0, None, "Interested-party", CTEE + "（未具名）", "—")
inp("b_eps27a", "券商 FY2027 EPS（凱基 08-13）", "NT$", None, 31.75, None, "Interested-party", KGI, "—")
inp("b_eps27b", "券商 FY2027 EPS（法人，10-08）", "NT$", None, 32.09, None, "Interested-party", CTEE + "（未具名；目標價 400＝12.5x）", "FTNN 2026-10-09 同值")
inp("b_gm27", "券商 FY2027 毛利率（凱基）", "%", None, 0.047, None, "Interested-party", KGI, "營益率 3.2%")
inp("rk_sh", "三家 ODM 中廣達的 NVL72 櫃數份額（8 月調查）", "%", None, "=2525/(2600+2525+2350)", None, "Interested-party", "經濟日報 2026-09-10 引述美系券商（網址見 Sources）", "單月、三家合計（不含 Dell、Supermicro 等）；第二個量化來源：未找到")
inp("pe_fair_andy", "Andy 先驗：合理 P/E", "x", None, 15, None, "Assumed", "先驗_Andy_廣達_20261024 Q9b（Drive；不在本 repo）", "只作對照")

sec("H. 價格與估值")
inp("px", "股價（2026-10-22 收盤）", "NT$", None, 322, None, "Interested-party", "Yahoo 股市 2382（10-22 收盤 322、前收 328.5）", "財報狗：10-13 收盤 323")
inp("pe_ttm", "近 12 月 P/E（10-22）", "x", None, 14.31, None, "Interested-party", "Yahoo 股市 2382", "—")

sec("I. 目標價情境權重（Val 第 0 節用；Andy 可改）")
for s_, w_, nm_ in (("A", 0.3, "順利"), ("B", 0.5, "平平"), ("C", 0.2, "不如預期")):
    inp(f"tp_w{s_}", f"目標價情境權重：情境 {s_}（{nm_}）", "權重", None, w_, None, "Assumed",
        "Andy 可改（2026-10-25 目標價規格預設 A／B／C＝0.3／0.5／0.2）", "三個權重合計應為 1；v0.2 起用於 DCF 與 Comparable 的情境加權（Val 第 0 節）")

# ---------------- v0.2 設定與新增輸入（附加在 In 既有輸入之後，不位移既有列） ----------------
V2 = {'name': '廣達',
 'unit': 'NT$B',
 'shunit': '十億股',
 'cur': 'NT$',
 'shg': 'sh_g',
 'epsfmt': '0.00',
 'pxfmt': '#,##0',
 'keys': {'rev': 'rev', 'ni': 'ni', 'sh': 'sh_', 'eps': 'eps'},
 'roe': ('=AVERAGE(0.2186,0.2842,0.3119)', 'stockanalysis 比率頁（Return on Equity）：FY2023 21.86%、FY2024 28.42%、FY2025 31.19%', '單一來源；近 12 月 39.48%（同站）'),
 'cons_tp': (427.84, 'stockanalysis 個股頁（S&P Global）19 位分析師平均 427.84，頁面時間 2026-10-14', '第二來源：未找到'),
 'cons_feps': ('=323/11.31', 'stockanalysis 個股頁 2026-10-14：價格 323 ÷ Forward PE 11.31'),
 'peers': [('緯創（Wistron）', '3231.TW', 9.93, '2026-09-29'),
           ('緯穎（Wiwynn）', '6669.TW', 13.87, '2026-10-02'),
           ('鴻海（Hon Hai）', '2317.TW', 11.37, '2026-10-19'),
           ('英業達（Inventec）', '2356.TW', 18.0, '2026-08-14'),
           ('Super Micro', 'SMCI', 7.06, '2026-10-23')],
 'peer_note': 'AI 伺服器 ODM 同業：緯創、緯穎、鴻海、英業達、Super Micro（券商常用比較組）。工業富聯（601138.SS）為 A 股，依規則不列。非 A 股公司的同業名單不放中國 A 股（A 股有結構性溢價；Andy 2026-10-25 11:53）',
 'fwd_w': (0.16666666666666666, 0.8333333333333334),
 'fwd_note': '約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）',
 'pe10': None,
 'px_note': '現價＝In 頁 px（2026-10-22 收盤 322；10-23 收盤未找到）',
 'method': 'DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年平均約 27%（再投資率約 15%），FY2027–29 驅動因子、FY2030–32 收斂法＋永續。Comparable：緯創、緯穎、鴻海、英業達、Super Micro 前瞻 P/E 中位數；部分價格為 8–10 月初，Super Micro 低倍數（約 '
           '7x）反映會計與治理折價，偏誤方向不明確。情境對應：A＝順利、B＝平平、C＝不如預期（原顆數情境）。',
 'readme_val': '估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），FY2027–FY2032 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，推到 12 個月後）＋0.5 × Comparable 目標價（緯創、緯穎、鴻海、英業達、Super Micro 前瞻 P/E 中位數 × '
               '12 個月後前瞻 EPS）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解（EPS 因素、倍數因素）、DCF 價格帶、十年 P/E 分位檢查（只作檢查）。',
 'readme_add': ['v0.2（2026-10-25）：(1) PL 頁最下方新增延伸表，三情境延伸到 FY2032；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt；淨利率由 FY2029 收斂到 A＝FY2029、B＝FY2029 '
                '與十年平均的中點、C＝十年平均；十年平均淨利率未找到，依工作單以 FY2029 × 0.8 替代，Assumed；PL 延伸表灰底）。(2) Val 第 0 節目標價＝DCF 與 Comparable 加權（tp_w_dcf＝0.5）；v0.1.1 的「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留，不再是目標價。(3) Sens '
                '新增目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.4／0.8）。(4) Recon 新增目標價對帳。(5) 現價：10-23 收盤未找到（Yahoo 股市、財報狗於 2026-10-25 仍為 10-22 或更早；markets/invest_map/data/prices.csv 亦為 10-22），依工作單沿用 '
                '2026-10-22 收盤 322。 新增內容一律放在各頁既有內容之後，FY2025–FY2029 既有數值不變。',
                '情境對應：A＝順利（顆數基準）、B＝平平（CY2028 顆數 ×0.8、CY2029 ×0.85）、C＝不如預期（CY2028–29 ×0.6），即原本的三個顆數情境。三情境在 FY2028 前完全相同（只差顆數係數，從 CY2028 起），DCF 三情境仍照算，差異來自 FY2028 之後（含 FY2030 起的收斂法）。',
                '先驗：v0.1 的 Andy 合理 P/E 對照保留不動，不進目標價；本版未新增先驗。']}
sec("J. 目標價（DCF 與 Comparable 加權）與收斂法參數（v0.2 新增；黃底＝Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（DCF 永續；情境 B 的 FY2032 營收年增收斂目標）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "Sens 頁 ±1pt（只動 DCF，PL 收斂目標不隨之變動）")
inp("tp_roe", "ROE（近三年平均）", "%", None, V2["roe"][0], None, "Derived", V2["roe"][1], V2["roe"][2])
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 改為 0.5：DCF／Comparable 各半）", "Sens 頁 0.4／0.8")
inp("fade_gA", "情境 A（順利）：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂，FY2032 達到")
inp("fade_gB", "情境 B（平平）：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C（不如預期）：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("nim10_f", "十年平均淨利率的替代：最後一年（FY2029）淨利率 ×（本值）", "x", None, 0.8, None, "Assumed", "十年（2016–2025）平均淨利率：未找到（免費資料商只到近五年）；依工作單以最後一年 × 0.8 替代", "淨利率收斂目標：A＝FY2029、B＝FY2029 與替代值的中點、C＝替代值（PL 延伸表）")
inp("cons_tp", "市場共識目標價（賣方平均）", V2["cur"], None, V2["cons_tp"][0], None, "Interested-party", V2["cons_tp"][1], V2["cons_tp"][2])
inp("cons_feps", "共識前瞻 12 月 EPS（＝價格 ÷ 前瞻 P/E）", V2["cur"], None, V2["cons_feps"][0], None, "Interested-party", V2["cons_feps"][1], "賣方共識，只作對帳與差距拆解，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣；股數為十億股）。藍字＝輸入；黑字＝公式。實績取自廣達法說摘要、財報新聞與月營收"
Q["A2"] = "FY2026 錨點＝1H26 實績＋3Q26 營收實績（月營收）＋3Q、4Q 假設（毛利率、營業費用率、淨利 ÷ 營業利益；4Q 營收＝3Q × (1＋季增)）。FY2025 與 1H26 毛利由公司揭露的毛利率 × 營收推得；FY2025 營業利益同。業外與稅前：未找到"
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（營收實績）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：" + UAN25 + "；1H26：" + FUGLE
S2 = "3Q26 營收：月營收公告（財報狗 2026-10-08）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "營收", {"B": 2123.689, "C": 1845.8, "D": "=q3_rev", "E": "=D4*(1+q4g)"}, tag="Verified／Interested-party（4Q 季增）", s2="FY2025 2,123.689（凱基）；1H26 1.8458 兆（公司）")
qrow(5, "毛利", {"B": "=B4*0.0698", "C": "=C4*0.0492", "D": "=D4*q_gm", "E": "=E4*q_gm"}, tag="Derived（毛利率 × 營收）／Interested-party", s2="FY2025 毛利率 6.98%；1H26 4.92%（公司）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(7, "營業費用", {"B": "=B5-B8", "C": "=C5-C8", "D": "=D4*q_opx", "E": "=E4*q_opx"}, tag="Derived（毛利 − 營業利益）／Assumed")
qrow(8, "營業利益", {"B": "=B4*0.0412", "C": 57.3, "D": "=D5-D7", "E": "=E5-E7"}, tag="Derived／Verified", s2="FY2025 營益率 4.12%；1H26 573 億、營益率 3.1%（公司）")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(10, "歸屬母公司淨利", {"B": 74.98, "C": 49.84, "D": "=D8*q_nio", "E": "=E8*q_nio"}, tag="Verified／Derived（3Q、4Q）", s2="2Q26 28.65、2Q26 EPS 7.43")
qrow(11, "淨利 ÷ 營業利益（含業外、稅、非控制權益）", {c: f"={c}10/{c}8" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(12, "股數（十億股）", {"B": "=74.98/19.45", "C": "=q_sh", "D": "=q_sh", "E": "=q_sh"}, ann="=q_sh", tag="Derived", s2="淨利 ÷ EPS")
qrow(13, "EPS", {c: f"={c}10/{c}12" for c in "BCDE"}, fmt="0.00", ann="=F10/F12", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(14, "  公司揭露 EPS", {"B": 19.45, "C": 12.93}, fmt="0.00", ann="", tag="Verified", s2="2Q26 7.43")
qrow(15, "AI 伺服器營收（＝營收 × 伺服器占比 × AI 占伺服器）", {"B": "=B4-F16/(1+g_nai26)"}, ann="=F4*srv_26*ais_26", tag="Derived／Interested-party", s1="公司比例（In 頁 C 節）；FY2025＝營收 − FY2026 非 AI ÷（1＋非 AI 年增）", s2="各季與年度 AI 營收金額未揭露：未找到")
qrow(16, "非 AI 營收（一般伺服器、筆電、其他）", {"B": "=F16/(1+g_nai26)"}, ann="=F4-F15", tag="Derived／Assumed（FY2025）", s1="—", s2="—")
Q.cell(18, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(19, 1, "FY2026 營收年增"); Q["B19"] = "=F4/B4-1"
Q.cell(20, 1, "FY2026 營業費用年增"); Q["B20"] = "=F7/B7-1"
Q.cell(21, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B21"] = "=B20/B19"
Q.cell(22, 1, "FY2026 營收 ÷ 凱基（08-13，3Q 前）− 1"); Q["B22"] = "=F4/b_rev26-1"
Q.cell(23, 1, "FY2026 EPS ÷ 凱基 − 1"); Q["B23"] = "=F13/b_eps26-1"
Q.cell(24, 1, "AI 伺服器營收 FY2026 年增（隱含；公司說三位數）"); Q["B24"] = "=F15/B15-1"
for rr in range(19, 25): Q[f"B{rr}"].number_format = "0.0%" if rr != 21 else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 10), ("sh", 12), ("eps", 13), ("ai", 15), ("nai", 16)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$F${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 90; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", rel="rel", cs="cs", keep="keep", vpg="vpg", mf="1", gn="g_nai", d="d_ai", vrd="vrd", opxa="0", nio="nio_f")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 部署與整櫃組裝市場（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029）")
    row("itc", "AI IT 資本支出 × 情境顆數係數", "US$B", {c_(y): f"=itc_{y % 100}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("itpg", "每 GW IT 資本支出", "US$B", {"D": "=itpg_gb*(1-vr26)+itpg_vr*vr26", **{c_(y): "=itpg_vr" for y in BU[1:]}})
    row("gw", "部署 GW（IT 電力）", "GW", {c_(y): f"={c_(y)}{R['itc']}/{c_(y)}{R['itpg']}" for y in BU}, fmt="0.0")
    row("vpg", "每 GW ODM 整櫃價值", "US$B/GW", {"D": "=(racks_gb*odm_gb*(1-vr26)+racks_vr*odm_vr*vr26)/1000", "E": "=racks_vr*odm_vr/1000",
                                                  **{c_(y): f"=$E{r}*{P['vpg']}_{y % 100}" for y in BU[2:]}}, fmt="0.00",
        note="架數 × ODM 每櫃價格；含 ASIC 機架（Assumed：每 GW 價值相同）")
    row("mku", "AI 伺服器組裝市場（ODM 價，美元）", "US$B", {c_(y): f"={c_(y)}{R['gw']}*{c_(y)}{R['vpg']}" for y in BU}, fmt="#,##0.0")
    row("mkn", "AI 伺服器組裝市場（新台幣）", "NT$B", {c_(y): f"={c_(y)}{R['mku']}*fx" for y in BU}, fmt="#,##0.0")
    row("sh0", "廣達有效份額（FY2026 校準）＝ AI 伺服器營收 ÷ 市場", "%", {"D": f"=a26_ai/D{R['mkn']}"}, fmt="0.0%",
        note="Derived：分子為公司伺服器占比 × AI 占比 × 營收；分母含 ASIC、HGX 機架，所以低於整櫃櫃數份額")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "廣達有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("aiv", "AI 伺服器毛額價值（買賣模式下的營收）", "NT$B", {c_(y): f"={c_(y)}{R['mkn']}*{c_(y)}{R['sh']}" for y in BU}, fmt="#,##0.0")
    row("cs", "客供比例", "%", {"D": "=0", **{c_(y): f"={P['cs']}_{y % 100}" for y in BU[1:]}}, fmt="0.0%",
        note="FY2026 的少量客供已含在錨點與校準份額內，記為 0")
    row("ai", "AI 伺服器營收", "NT$B", {"C": "=a25_ai", **{c_(y): f"={c_(y)}{R['aiv']}*(1-{c_(y)}{R['cs']})" for y in BU}}, fmt="#,##0.0")
    sec_("B. 非 AI 與合計")
    row("nai", "非 AI 營收（一般伺服器、筆電、其他）", "NT$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']}{y % 100})" for y in BU[1:]}}, fmt="#,##0.0")
    row("rev", "營收", "NT$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.0")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI 伺服器 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("pgw", "每 GW 廣達 AI 伺服器營收", "US$B/GW", {c_(y): f"={c_(y)}{R['ai']}/fx/{c_(y)}{R['gw']}" for y in BU}, fmt="0.00")
    row("rk", "隱含廣達出貨櫃數（毛額價值 ÷ ODM 每櫃價格；NVL72 等值）", "千櫃",
        {"D": f"=D{R['aiv']}/fx/(odm_gb*(1-vr26)+odm_vr*vr26)", **{c_(y): f"={c_(y)}{R['aiv']}/fx/odm_vr" for y in BU[1:]}}, fmt="#,##0.0",
        note="十億美元 ÷ 百萬美元 ＝ 千櫃；FY2028 起若每 GW 價值改變，每櫃價格仍以 VR200 換算（參考）")
    sec_("C. 損益")
    row("gmn", "其他營收毛利率 ＝ 錨點毛利率 − 差距 × AI 占比（FY2026 校準）", "%", {"D": f"=a26_gp/D{R['rev']}-{P['d']}*D{R['aish']}"}, fmt="0.00%",
        note="Derived：讓 FY2026 毛利率等於錨點；AI 毛利率（毛額基礎）＝其他 + 差距（差距為負）")
    row("gai", "AI 伺服器毛利 ＝ AI 毛利率 × 毛額價值 ×（1 − 客供 ×（1 − 保留比例））", "NT$B",
        {c_(y): f"=($D${R['gmn']}+{P['d']}" + ("" if y == 2026 else f"+{P['vrd']}") + f")*{c_(y)}{R['aiv']}*(1-{c_(y)}{R['cs']}*(1-{P['keep']}))" for y in BU}, fmt="#,##0.00",
        note="客供部分不計營收，但大部分毛利金額保留（keep）；FY2027 起加 VR 稀釋 vrd")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"=$D${R['gmn']}*{c_(y)}{R['nai']}+{c_(y)}{R['gai']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("gma", "AI 伺服器毛利率（營收基礎）", "%", {c_(y): f"={c_(y)}{R['gai']}/{c_(y)}{R['ai']}" for y in BU}, fmt="0.00%",
        note="客供比例上升 → 營收基礎毛利率上升")
    row("opb", "營業費用（基準路徑）", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in BU[1:]}}, fmt="#,##0.0")
    row("opx", "營業費用", "NT$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("oi", "營業利益", "NT$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.0")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("nio", "淨利 ÷ 營業利益", "%", {"C": "=a25_ni/a25_oi", "D": "=a26_ni/a26_oi", **{c_(y): f"={P['nio']}" for y in BU[1:]}}, fmt="0.0%")
    row("ni", "歸屬母公司淨利", "NT$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"={c_(y)}{R['oi']}*{c_(y)}{R['nio']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.000")
    row("eps", "EPS", "NT$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 廣達損益（十億新台幣；EPS 新台幣；歸屬母公司）。FY2025A 實績、FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 伺服器營收", "NT$B", "#,##0.0"),
                                ("nai", "  非 AI 營收", "NT$B", "#,##0.0"), ("aish", "  AI 伺服器 ÷ 營收", "%", "0.0%"), ("gp", "毛利", "NT$B", "#,##0.0"),
                                ("gm", "毛利率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"),
                                ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("eps", "EPS", "NT$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收 − 錨點；FY2026 模型 AI 伺服器 − 錨點；FY2026 毛利 − 錨點").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：AI 伺服器差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['ai']}-a26_ai"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（NT$B；以 FY2027 公式套用 FY2026）"); PL[f"D{sr}"] = f"=D{PB['A']['gmn']}*D{PB['A']['nai']}+D{PB['A']['gai']}-a26_gp"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：組裝市場、份額、客供（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — AI 伺服器組裝市場、廣達份額、客供比例（本專案權威，規則第 13 條）。部署 GW＝AI IT 資本支出（markets/network）÷ 每 GW IT 資本支出（Tokenomics，經橋接表）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("gw", "部署 GW（情境 A）", "GW", lambda y: f"=PL!{FC[y]}{A_['gw']}", "0.0", "AI IT 資本支出 ÷ 每 GW IT 資本支出"),
        ("vpg", "每 GW ODM 整櫃價值", "US$B/GW", lambda y: f"=PL!{FC[y]}{A_['vpg']}", "0.00", "架數 × ODM 每櫃價格（GB300 4.0、VR200 6.75 百萬美元）"),
        ("mku", "AI 伺服器組裝市場（ODM 價）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}", "#,##0.0", "含 ASIC 機架（假設每 GW 價值相同）"),
        ("sh", "廣達有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以公司伺服器與 AI 占比校準"),
        ("aiv", "廣達 AI 伺服器毛額價值", "US$B", lambda y: f"=PL!{FC[y]}{A_['aiv']}/fx", "#,##0.0", "買賣模式下的營收"),
        ("cs", "客供比例", "%", lambda y: f"=PL!{FC[y]}{A_['cs']}", "0.0%", "Assumed（FY2027 起）"),
        ("ai", "廣達 AI 伺服器營收", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx", "#,##0.0", ""),
        ("rk", "隱含廣達出貨櫃數（NVL72 等值）", "千櫃", lambda y: f"=PL!{FC[y]}{A_['rk']}", "#,##0.0", "毛額價值 ÷ ODM 每櫃價格"),
        ("pgw", "每 GW 廣達 AI 伺服器營收", "US$B/GW", lambda y: f"=PL!{FC[y]}{A_['pgw']}", "0.00", "先驗比對推導 2026 約 3.4–4.2"),
        ("gma", "廣達 AI 伺服器毛利率（營收基礎）", "%", lambda y: f"=PL!{FC[y]}{A_['gma']}", "0.00%", "客供上升 → 毛利率上升、毛利金額大致不變"),
        ("oth", "其他廠商（鴻海、緯創、英業達、Dell、Supermicro 等）＝ 市場 − 廣達毛額", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}-PL!{FC[y]}{A_['aiv']}/fx", "#,##0.0", "")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：FY2026 有效份額遠低於整櫃櫃數份額（三家中約三分之一），因為本模型分母是全部 AI IT 資本支出換算的組裝市場（含 ASIC、HGX 與其他組裝廠），且公司部分出貨為運算托盤而非整櫃。客供比例與份額相對值是 Assumed，是最需要每季檢查的輸入。").font = BOLD
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」與十年 P/E 規則值（年底落後 P/E，基準為 FY2027 EPS）"
V["A2"] = "循環位置：廣達不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 NVIDIA、Arista、台達電；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-22 收盤）"); V["B6"] = "=px"; V["B6"].number_format = "#,##0"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "未來 12 月 EPS（模型 A）≈ 0.2 × FY2026＋0.8 × FY2027"); V["B11"] = f"=0.2*PL!D{SUMR[('A', 'eps')]}+0.8*PL!E{SUMR[('A', 'eps')]}"; V["B11"].number_format = "0.00"
V.cell(12, 1, "  現價 ÷ 未來 12 月 EPS（模型 A）"); V["B12"] = "=px/B11"; V["B12"].number_format = "0.0"
V.cell(13, 1, "近 12 月 P/E（Yahoo，10-22）"); V["B13"] = "=pe_ttm"; V["B13"].number_format = "0.0"
V["E8"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"; V["E12"] = "Andy 先驗 Q9a 答 20x（未來 12 月）"
V.cell(15, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, 16, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 17):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(21, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 22, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 23 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{17 + j}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(26, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, 27, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，C）", "對照：Andy 合理 P/E", "對照：十年中位數"])
vr = 28
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, "=pe_fair_andy").number_format = "0.0"; V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。")
vr += 2
V.cell(vr, 1, "5. 對照：以 Andy 合理 P/E 15x 計算的隱含價（只作對照，非規則值）").font = BOLD; vr += 1
hdr(V, vr, ["情境", "FY2027 EPS", "× 15x", "相對現價", "FY2028 EPS", "× 15x", "相對現價"]); vr += 1
IMP = {}
for s in SC:
    V.cell(vr, 1, s)
    V.cell(vr, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 3, f"=B{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 4, f"=C{vr}/px-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 6, f"=E{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
    IMP[s] = vr; vr += 1
assert vr < 50, vr
PE_CMC = [9.26, 10.5, 8.95, 11.1, 9.23, 8.62, 8.30, 20.2, 17.7, 13.5]       # companiesmarketcap 年底 P/E 2016–2025
PE_GI = ["=60.3/3.93", "=61.9/3.73", "=52.7/3.92", "=64.3/4.14", "=80.9/6.57", "=94.7/8.73", "=72.3/7.51", "=224.5/10.29", "=287/15.49", "=272/19.45"]  # Goodinfo 年底收盤 ÷ EPS
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2016–2025 共 10 年；價格可能為還原股價）"},
    {"label": "Goodinfo（年底收盤 ÷ EPS）", "vals": PE_GI, "tag": "Derived（資料商收盤價與 EPS；未還原除息）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。兩來源 2016–2021 年相差約 3–6 個倍數點（companiesmarketcap 疑用還原股價，數字較低）；Goodinfo 多年加權平均 P/E 區間 12.02–20.7 只作量級核對。2016–2022 年廣達為筆電代工為主的低成長公司（約 8–17x），2023 年起 AI 伺服器重新定價，規則值可能低估 AI 期的合理倍數。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# ---------------- Val 第 0 節：目標價與上檔空間（v0.1.1 新增；放在既有空白列，不位移既有列） ----------------
TP0 = 41; TPR = {}
assert all(V.cell(rw, c).value is None for rw in range(TP0 - 1, TP0 + 9) for c in range(1, 9)), TP0
assert V["A3"].value is None
V["A3"] = f"目標價與上檔空間：見第 0 節（本頁第 {TP0} 列；v0.1.1 新增，為了不位移既有列而放在空白列）"
V.cell(TP0, 1, "舊法：規則值情境加權（參考；v0.1.1 的「0. 目標價與上檔空間」改名，不再是目標價；權重為 In 頁 tp_wA／tp_wB／tp_wC）").font = BOLD
for k_, (key_, lab_) in enumerate((("tp", "舊法隱含價（情境加權，參考）"), ("up", "舊法隱含價相對現價（÷ 現價 − 1）"),
        ("A", "情境 A 隱含價（樂觀）"), ("B", "情境 B 隱含價（基準）"), ("C", "情境 C 隱含價（下檔）"),
        ("w", "情境權重 A／B／C"), ("m", "舊法說明")), TP0 + 1):
    V.cell(k_, 1, lab_); TPR[key_] = k_
V.cell(TPR["tp"], 2, f"=tp_wA*B{TPR['A']}+tp_wB*B{TPR['B']}+tp_wC*B{TPR['C']}").number_format = "#,##0"
V.cell(TPR["tp"], 5, "＝權重 A × 情境 A 隱含價＋權重 B × 情境 B 隱含價＋權重 C × 情境 C 隱含價（NT$）")
V.cell(TPR["up"], 2, f"=B{TPR['tp']}/px-1").number_format = "0.0%"
V.cell(TPR["up"], 5, "相對「" + V["A6"].value + "」")
for s_ in SC:
    V.cell(TPR[s_], 2, f"=E{VAL[s_]}").number_format = "#,##0"
    V.cell(TPR[s_], 5, f"＝規則值隱含股價 E{VAL[s_]}（{SCN_T[s_]}）")
for j_, s_ in enumerate(SC):
    V.cell(TPR["w"], 2 + j_, f"=tp_w{s_}").number_format = "0%"
V.cell(TPR["w"], 5, "In 頁輸入（Assumed，Andy 可改）：B 欄＝情境 A、C 欄＝情境 B、D 欄＝情境 C")
V.cell(TPR["m"], 2, "P/E 規則值：各情境 FY2027 EPS × 2016–2025 年底 P/E 依循環位置取分位（見頂＝25 分位、中段＝中位數、谷底＝75 分位；companiesmarketcap 與 Goodinfo 兩來源分別取分位再平均）。隱含價約為 FY2027 年底價、未折現，作 12 個月目標價的近似；情境 A／B／C 即本模型三個顆數情境（A 最高、C 最低）。偏誤：十年倍數含 2016–2022 年筆電代工期低倍數（約 8–17x），目標價可能偏低。")
for c_ in range(1, 6):
    V.cell(TPR["tp"], c_).fill = KEY; V.cell(TPR["up"], c_).fill = KEY
V.cell(TPR["tp"], 1).font = BOLD
defname("old_rule_px", f"Val!$B${TPR['tp']}"); defname("old_rule_up", f"Val!$B${TPR['up']}")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額、其他營收毛利率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 伺服器營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 伺服器營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 伺服器營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("份額流失（低）", "份額相對值 2027 0.85、2028 0.75、2029 0.70（鴻海、緯創搶單）", dict(rel="RLO")),
         ("份額上升（高）", "份額相對值 1.1（Neocloud、ASIC 專案）", dict(rel="RHI")),
         ("客供 0（維持買賣）", "客供比例 2027–29 皆 0", dict(cs="CSLO")), ("客供 高", "客供比例 25%／40%／50%", dict(cs="CSHI")),
         ("客供保留毛利 0.7", "客供部分只保留 70% 毛利", dict(keep="keep_lo")), ("客供保留毛利 1.0", "毛利金額完全保留", dict(keep="keep_hi")),
         ("每 GW 價值 低", "2028–29 每 GW ODM 價值 ×0.85（Rubin Ultra 降價或客供更深）", dict(vpg="VPLO")), ("每 GW 價值 高", "×1.2", dict(vpg="VPHI")),
         ("AI 部署 ×0.8", "FY2027 起 AI IT 資本支出 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("非 AI 成長 低", "2027–29 −5%", dict(gn="GNLO")), ("非 AI 成長 高", "2027 +8%、2028–29 +6%", dict(gn="GNHI")),
         ("AI 毛利率差距 −3.5pt", "AI 毛利率（毛額基礎）比其他低 3.5 個百分點", dict(d="d_ai_lo")), ("AI 毛利率差距 −0.5pt", "只低 0.5 個百分點", dict(d="d_ai_hi")),
         ("VR 稀釋 −0.6pt", "FY2027 起 AI 毛利率（毛額基礎）−0.6 個百分點", dict(vrd="vrd_lo")), ("VR 不稀釋", "FY2027 起不變", dict(vrd="vrd_hi")),
         ("營業費用率 +0.2pt", "營業費用額外 +0.2% × 營收", dict(opxa="opx_adj_hi")),
         ("淨利 ÷ 營業利益 0.75", "FY2027 起 0.75（券商隱含約 0.73–0.77）", dict(nio="nio_f_lo"))]
for nm, base, suf in (("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("CSLO", "cs", "_lo"), ("CSHI", "cs", "_hi"), ("VPLO", "vpg", "_lo"), ("VPHI", "vpg", "_hi"), ("GNLO", "g_nai", "_lo"), ("GNHI", "g_nai", "_hi")):
    for y in (27, 28, 29):
        if base == "vpg" and y == 27: continue
        src = f"{base}{y}{suf}" if base == "g_nai" else f"{base}_{y}{suf}"
        target = f"{nm}{y}" if base == "g_nai" else f"{nm}_{y}"
        wb.defined_names[target] = DefinedName(target, attr_text=wb.defined_names[src].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "ai", 2027, "#,##0.0"), ("E", "gm", 2027, "0.00%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "ai", 2028, "#,##0.0"), ("J", "gm", 2028, "0.00%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數情境 B、C 的差異從 FY2028 起。每個案例的 FY2026 有效份額與其他營收毛利率依該案例重算，所以 FY2026 恆等於錨點。「客供」案例同時改變營收與毛利率，EPS 只受保留毛利比例影響：客供比例高、保留比例低時 EPS 才明顯下降。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 券商估計與櫃數份額。只作事後對照，不為對上數字而調參數。券商多未具名或單一來源；2027 共識（付費）未找到"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 券商估計（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 凱基（08-13，3Q 前）", "=b_rev26", pl("rev", 2026), "凱基 2026-08-13", "3Q 只達法人預期 92%", fmt="#,##0.0")
rec("FY2026 EPS vs 凱基", "=b_eps26", pl("eps", 2026), "凱基 2026-08-13", "")
rec("FY2027 營收 vs 凱基", "=b_rev27a", pl("rev", 2027), "凱基 2026-08-13", "", fmt="#,##0.0")
rec("FY2027 營收 vs 法人（10-08）", "=b_rev27b", pl("rev", 2027), "工商時報 2026-10-08", "", fmt="#,##0.0")
rec("FY2027 EPS vs 凱基", "=b_eps27a", pl("eps", 2027), "凱基 2026-08-13", "")
rec("FY2027 EPS vs 法人（10-08）", "=b_eps27b", pl("eps", 2027), "工商時報 2026-10-08", "兩家區間 31.75–32.09")
rec("FY2027 毛利率 vs 凱基（百分點）", "=b_gm27", pl("gm", 2027), "凱基 2026-08-13", "模型含客供 10%，毛利率較高", fmt="0.00%", pt=True)
rec("FY2027 營收年增（模型）", None, pl("g", 2027), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "B. 份額與每 GW（先驗比對的機制檢查）").font = BOLD; rr_ += 1
rec("FY2026 有效份額 vs 三家 ODM 中的 NVL72 櫃數份額", "=rk_sh", lambda s: f"=PL!D{PB[s]['sh0']}", "經濟日報 2026-09-10（美系券商）", "口徑不同：本模型分母為全部 AI 組裝市場（含 ASIC、HGX、其他廠）", fmt="0.0%", pt=True)
rec("FY2026 每 GW 廣達 AI 伺服器營收（US$B）", None, pl("pgw", 2026), "—", "先驗比對推導約 3.4–4.2", fmt="0.00")
rec("FY2026 隱含廣達出貨櫃數（千櫃，NVL72 等值）", None, pl("rk", 2026), "—", "三家調查單月約 2.5 千櫃 × 12 ≈ 30 千櫃（Interested-party；含運算托盤口徑不明）", fmt="#,##0.0")
rec("FY2027 AI 伺服器 ÷ 營收", None, pl("aish", 2027), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("其他營收毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "AI 毛利率（毛額基礎）＝此值＋d_ai", fmt="0.00%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "客供使毛利率上升", fmt="0.00%", pt=True)
rec("4Q26 EPS（模型）", None, lambda s: "=Q!E13", "—", "", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代「AI IT 資本支出 ÷ 每 GW IT 資本支出」"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文只經法說摘要與新聞轉述，原始簡報待人工複核；券商與資料商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司（摘要）：富果 2Q26 法說摘要 2026-08-13（1H26 損益、伺服器與 AI 占比、客供、產能、資本支出）", FUGLE),
        ("公司（摘要）：優分析 2Q26 法說（AI 占伺服器 75–80%、全年約 80%、毛利率、能見度至 2028）", UAN),
        ("公司（摘要）：Taipei Times 2026-05-15（1Q26 毛利率 4.78%、客供轉換協商）", "https://www.taipeitimes.com/News/biz/archives/2026/05/15/2003857345"),
        ("二手：優分析 2025 年財報（毛利率 6.98%、營益率 4.12%、淨利 749.8 億、EPS 19.45）", UAN25),
        ("二手：財報狗 2026-10-08（3Q26 營收 1.1155 兆、前三季 2.9614 兆、VR 時程、筆電占比）", SD1008),
        ("券商：凱基 2026-08-13（2025A–2027F 營收、EPS、毛利率、營益率）", KGI),
        ("券商：工商時報 2026-10-08（法人 2027 營收 5.3 兆、EPS 32.09、GB300 與 VR 每櫃價格）", CTEE),
        ("券商：FTNN 2026-10-09（同上；Neocloud 繞過 OEM）", "https://www.ftnn.com.tw/news/585951"),
        ("券商：經濟日報 2026-09-10（美系券商 NVL72 櫃數調查）", UDN0910),
        ("資料商：Yahoo 股市（10-22 收盤、近 12 月 P/E）", "https://tw.stock.yahoo.com/quote/2382.TW"),
        ("資料商：companiesmarketcap 年底 P/E", "https://companiesmarketcap.com/quanta-computer/pe-ratio/"),
        ("資料商：Goodinfo 歷年經營績效（年底收盤、EPS）", "https://goodinfo.tw/tw/StockBzPerformance.asp?STOCK_ID=2382&RPT_CAT=M_YEAR"),
        ("市場層：markets/network、markets/chip_bridge、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：3Q26 法說（11 月中）；客戶占比（年報 10% 以上客戶，MOPS）；業外與稅前利益；客供比例的量化；前瞻 P/E 十年歷史（Bloomberg）；券商共識（付費）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("廣達電腦（2382）收支與估值模型 v0.1.1（AI 半導體 Project：伺服器組裝次產業樣板）　2026-10-25", True),
 ("版本：v0.1.1（2026-10-25：Val 新增第 0 節「目標價與上檔空間」、In 新增情境權重 tp_wA／tp_wB／tp_wC；其餘數值與 v0.1 相同）。v0.1（初版）：AI 伺服器組裝市場、廣達有效份額、ODM 每櫃價格與客供比例在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）；鴻海、緯創之後讀這條組裝線。", False),
 ("模型命題：廣達的 AI 伺服器營收＝部署 GW（AI IT 資本支出 ÷ 每 GW IT 資本支出）× 每 GW ODM 整櫃價值 × 廣達有效份額 ×（1 − 客供比例）。組裝營收含轉手的 GPU，營收大、毛利率約 5%；客供比例上升時營收成長放慢、毛利率回升，毛利金額大致不變。市場約 10 倍未來 12 月 EPS。", False),
 ("核心問題：客供比例上升與 VR 每櫃價格上升之下，每櫃毛利金額能不能撐住 EPS 成長；約 10 倍未來 12 月 P/E 隱含什麼路徑。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029（三年）。金額：十億新台幣；整櫃價格與市場以美元計、以匯率換算。口徑：歸屬母公司淨利。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 伺服器營收：markets/network 的 AI IT 資本支出 × 情境顆數係數 ÷ 每 GW IT 資本支出（Tokenomics，經橋接表）＝部署 GW；× 每 GW 架數（Tokenomics）× ODM 每櫃價格（法人）＝組裝市場；× 廣達有效份額（FY2026 以公司伺服器與 AI 占比校準）× 份額相對值（Assumed）×（1 − 客供比例，Assumed）。", False),
 ("非 AI：FY2026 殘差 × 成長率（Assumed）。毛利率：AI 伺服器（毛額基礎）比其他營收低 d_ai（Assumed），FY2027 起再加 VR 稀釋 vrd（Assumed），其他毛利率為 FY2026 殘差（Derived）；客供部分不計營收、保留 keep 比例的毛利。營業費用：彈性 phi（Derived）。淨利：營業利益 × 淨利係數（業外與稅前未找到）。", False),
 ("校準：FY2026 ＝ 1H26 實績＋3Q26 營收實績＋假設（4Q 季增、毛利率、營業費用率）。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：現價 ÷ FY2027、FY2028 EPS；十年 P/E 規則值（兩個資料商來源）× FY2027 EPS；循環位置以本模型營收最高年判定（不套設備股見頂年）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文只經法說摘要與新聞；券商多未具名或單一來源。", False),
 ("", False),
 ("已知限制：(1) 客供比例與保留毛利比例是 Assumed，公司只說「比重仍低」；(2) 份額只有一個量化來源（單月櫃數調查），且有效份額的分母含 ASIC、HGX 與其他廠，與櫃數份額不同口徑；(3) 2028 年起每 GW ODM 價值（Rubin Ultra）是 Assumed；(4) 部署 GW 以 VR200 每 GW IT 資本支出換算，ASIC 機架每 GW 資本支出較低，GW 可能低估；(5) 業外與稅前利益未找到，以淨利 ÷ 營業利益涵蓋；(6) 十年 P/E 兩來源在 2016–2021 年差距大，且 AI 前為低倍數；(7) 3Q26 法說（11 月中）後更新錨點；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

V2["pe10"] = (list(zip(range(2016, 2026), [f"=({a}+{str(b).lstrip('=')})/2" for a, b in zip(PE_CMC, PE_GI)])), "companiesmarketcap 與 Goodinfo 兩來源平均；同規則值")
V2["old_row"] = TP0
# ================= v0.2（2026-10-25）：PL 延伸到 FY2032、Val 第 0 節目標價（DCF 與 Comparable 加權）、Sens 目標價參數 =================
# 全部放在各頁既有內容之後，不位移既有儲存格（FY2025–FY2029 數值與 v0.1.1 相同）
FADEFILL = PatternFill("solid", fgColor="E7E6E6")
XY = (2025, 2026, 2027, 2028, 2029, 2030, 2031, 2032); XC = {y: "CDEFGHIJ"[i] for i, y in enumerate(XY)}; FADE = (2030, 2031, 2032)
K2 = V2["keys"]
for rw in range(4, IN.max_row + 1):
    if IN.cell(rw, 1).value in ("tp_wA", "tp_wB", "tp_wC", "tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "nim10_f"):
        IN.cell(rw, 5).fill = KEY
# --- PL 延伸表（三情境到 FY2032） ---
xr = PL.max_row + 3; XT = xr
PL.cell(xr, 1, "延伸：三情境到 FY2032（v0.2 新增）。FY2025–FY2029 引用上方推導區塊；FY2030 起為收斂法（灰底），非驅動因子推導：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 線性收斂到收斂目標，FY2032 達到").font = BOLD
xr += 1; XR = {}; XH = None
for s in SC:
    E = PB[s]; R = {}
    PL.cell(xr, 1, f"延伸｜{SCN_T[s]}").font = BOLD; xr += 1
    hdr(PL, xr, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    XH = XH or xr; xr += 1
    for key, lab, unit in (("rev", "營收（延伸）", V2["unit"]), ("g", "  營收年增", "%"), ("ni", "淨利（延伸）", V2["unit"]), ("nim", "  淨利率", "%"),
                           ("sh", "股數（延伸）", V2["shunit"]), ("eps", "EPS（延伸）", V2["cur"]), ("epsg", "  EPS 年增", "%"),
                           ("tg", "收斂目標：FY2032 營收年增", "%"), ("tm", "收斂目標：FY2032 淨利率", "%")):
        PL.cell(xr, 1, lab); PL.cell(xr, 2, unit); R[key] = xr; xr += 1
    G = lambda k: f"$G${R[k]}"
    for y in XY:
        c = XC[y]; pc = XC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={c}{E[K2['rev']]}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            PL[f"{c}{R['ni']}"] = f"={c}{E[K2['ni']]}"
            if y > 2025 or K2.get("sh25", True): PL[f"{c}{R['sh']}"] = f"={c}{E[K2['sh']]}"
            PL[f"{c}{R['eps']}"] = f"={c}{E[K2['eps']]}"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"={G('g')}+($C${R['tg']}-{G('g')})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"={G('nim')}+($C${R['tm']}-{G('nim')})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['sh']}"] = f"={pc}{R['sh']}*(1+{V2['shg']})"
            PL[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh']}"
        if y <= 2029: PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
        if y > 2025: PL[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("rev", "#,##0.0"), ("g", "0.0%"), ("ni", "#,##0.0"), ("nim", "0.0%"), ("sh", "0.0000"), ("eps", V2["epsfmt"]), ("epsg", "0.0%")):
            PL[f"{c}{R[key]}"].number_format = fm
            if y in FADE: PL[f"{c}{R[key]}"].fill = FADEFILL
    PL[f"C{R['tg']}"] = f"=fade_g{s}"; PL[f"C{R['tg']}"].number_format = "0.0%"
    PL[f"C{R['tm']}"] = {"A": f"=G{R['nim']}", "B": f"=G{R['nim']}*(1+nim10_f)/2", "C": f"=G{R['nim']}*nim10_f"}[s]; PL[f"C{R['tm']}"].number_format = "0.0%"
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導（上游市場層只到 2029）")
    PL.cell(R["tg"], 11, {"A": "In 頁 fade_gA＝tp_g＋2pt", "B": "In 頁 fade_gB＝tp_g", "C": "In 頁 fade_gC＝tp_g−2pt"}[s])
    PL.cell(R["tm"], 11, {"A": "＝本情境 FY2029 淨利率", "B": "＝FY2029 淨利率與十年平均替代值（FY2029 × nim10_f）的中點", "C": "＝十年平均替代值（FY2029 × nim10_f；十年平均未找到）"}[s])
    PL.cell(R["sh"], 11, f"FY2030 起股數年變動＝{V2['shg']}")
    XR[s] = R; xr += 1
PL.column_dimensions["I"].width = 14; PL.column_dimensions["J"].width = 14; PL.column_dimensions["K"].width = 60
# --- Val 第 0 節：目標價 ---
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
t0 = V.max_row + 3; T0 = t0
V.cell(t0, 1, "0. 目標價（DCF 與 Comparable 加權；v0.2 新增，放在本頁既有內容之後以免位移既有儲存格；三情境欄 B／C／D＝A 順利／B 平平／C 不如預期）").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["h1"] = t0 + 2
r_ = t0 + 3
V.cell(r_, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{r_}"] = "PL 頁延伸表；FY2030 起為收斂法"; T["eps_h"] = r_; r_ += 1
EPSR = {}
for y in YRS6:
    tset(r_, f"eps{y}", f"  FY{y}", [f"=PL!{XC[y]}{XR[s]['eps']}" for s in SC], V2["epsfmt"]); EPSR[y] = r_; r_ += 1
V.cell(r_, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = r_; r_ += 1
tset(r_, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）"); RV = r_; r_ += 1
CFR = {}
for y in YRS6:
    tset(r_, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], V2["epsfmt"]); CFR[y] = r_; r_ += 1
def dcf_formula(c, r="tp_r", g="tp_g", reinv=None):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    reinv = reinv or f"Val!$B${RV}"
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
c1, c6 = CFR[2027], CFR[2032]
tset(r_, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], V2["pxfmt"],
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027"); r_ += 1
tset(r_, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{T['dcfv']}*(1+tp_r)" for c in "BCD"], V2["pxfmt"]); r_ += 1
tset(r_, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{T['dcf12']}+tp_wB*C{T['dcf12']}+tp_wC*D{T['dcf12']}", None, None], V2["pxfmt"], "Σ tp_wA／wB／wC × 情境價值 ×(1＋tp_r)；權重見 In 頁"); r_ += 2
V.cell(r_, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; r_ += 1
hdr(V, r_, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["ph"] = r_; r_ += 1
P1 = r_
for nm, tk, pe, dt in V2["peers"]:
    V.cell(r_, 1, nm); V.cell(r_, 2, tk); cc = V.cell(r_, 3, pe); cc.font = BLUE; cc.number_format = "0.0"
    V.cell(r_, 4, dt); V.cell(r_, 5, "stockanalysis.com 個股頁 Forward PE（分析師共識；價格日期如左）"); V.cell(r_, 6, "Interested-party（賣方共識）"); r_ += 1
P2 = r_ - 1; T["p1"], T["p2"] = P1, P2
tset(r_, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{P1}:C{P2})", None, None], "0.0", V2["peer_note"]); r_ += 1
w1, w2 = V2["fwd_w"]
tset(r_, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*({w1}*{c}{EPSR[2027]}+{w2}*{c}{EPSR[2028]})" for s, c in zip(SC, "BCD")), None, None], V2["epsfmt"], V2["fwd_note"]); r_ += 1
tset(r_, "comptp", "Comparable 目標價", [f"=B{T['pmed']}*B{T['feps']}", None, None], V2["pxfmt"], "同業中位數 × 上列 EPS"); r_ += 2
V.cell(r_, 1, "0.3 結論").font = BOLD; r_ += 1
hdr(V, r_, ["項目", "數值", "", "", "說明"]); T["h3"] = r_; r_ += 1
tset(r_, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], V2["pxfmt"], "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁"); r_ += 1
tset(r_, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], V2["pxfmt"], "引用 0.2"); r_ += 1
tset(r_, "cons", "市場共識目標價", ["=cons_tp", None, None], V2["pxfmt"], V2["cons_tp"][1]); r_ += 1
tset(r_, "upw", "上檔空間：加權目標價", [f"=B{T['tpw']}/px-1", None, None], "0.0%", V2["px_note"]); r_ += 1
tset(r_, "upc", "上檔空間：Comparable 目標價", [f"=B{T['comp2']}/px-1", None, None], "0.0%"); r_ += 1
tset(r_, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS 見 In 頁 cons_feps 與 Recon"); r_ += 1
tset(r_, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%"); r_ += 1
V[f"E{T['gmul']}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(r_, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）"); r_ += 1
PEROW = r_ + 3
PE10 = f"$B${PEROW + 1}:$K${PEROW + 1}"
tset(r_, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
     [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E 分布（下方列）的分位；低於最小值記 0%、高於最大值記 100%；只作檢查，不進目標價"); r_ += 1
V.cell(r_, 1, "目標價方法"); V[f"B{r_}"] = V2["method"]; T["m"] = r_; r_ += 1
V.cell(PEROW, 1, "十年年底 P/E（歷史分位檢查用；" + V2["pe10"][1] + "）").font = BOLD
for j, (yy, vv) in enumerate(V2["pe10"][0]):
    V.cell(PEROW, 2 + j, str(yy))
    cc = V.cell(PEROW + 1, 2 + j, vv); cc.font = BLUE; cc.number_format = "0.0"
V.cell(PEROW + 1, 1, "年底 P/E"); T["pe10h"] = PEROW; T["pe10"] = PEROW + 1
for k_ in ("dcftp", "comptp", "tpw"):
    V[f"B{T[k_]}"].fill = KEY
V.cell(T["tpw"], 1).font = BOLD
defname("tp_final", f"Val!$B${T['tpw']}"); defname("tp_dcf", f"Val!$B${T['dcftp']}"); defname("tp_comp", f"Val!$B${T['comptp']}")
V["A3"] = f"目標價：見第 0 節（本頁第 {T0} 列；v0.2 新增，DCF 與 Comparable 加權）。第 {V2['old_row']} 列「舊法：規則值情境加權（參考）」不再是目標價，只作參考"
for c in "IJK": V.column_dimensions[c].width = 10
# --- Sens：目標價參數 ---
TPS0 = SE.max_row + 3
SE.cell(TPS0, 1, "目標價參數（v0.2）：一次動一個，對加權目標價的影響（只改 DCF 參數；PL 收斂目標 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = V2["pxfmt"]
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "tp_r 10%、tp_g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "tp_r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "tp_r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "tp_g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "tp_g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.4", "tp_w_dcf 0.4", dict(w="0.4")), ("DCF 權重 0.8", "tp_w_dcf 0.8", dict(w="0.8"))]
for k_, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k_, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
# --- Recon：目標價對帳 ---
rq = RC.max_row + 2
RC.cell(rq, 1, "E. 目標價對帳（v0.2；模型欄為情境加權值，非情境 A）").font = BOLD; rq += 1
RECV2 = rq
for lab, ref, mdl, src in (("12 個月後的前瞻 EPS（情境加權）vs 共識前瞻 EPS", "=cons_feps", f"=Val!B{T['feps']}", V2["cons_feps"][1]),
                           ("目標價（DCF 與 Comparable 加權）vs 市場共識目標價", "=cons_tp", f"=Val!B{T['tpw']}", V2["cons_tp"][1]),
                           ("Comparable 目標價 vs 市場共識目標價", "=cons_tp", f"=Val!B{T['comptp']}", "同上")):
    RC.cell(rq, 1, lab); RC.cell(rq, 2, ref).number_format = V2["epsfmt"]; RC.cell(rq, 3, mdl).number_format = V2["epsfmt"]
    RC.cell(rq, 4, f"=C{rq}/B{rq}-1").number_format = "0.0%"; RC.cell(rq, 9, src); RC.cell(rq, 10, "共識只作對帳，不作校準"); rq += 1
# --- README ---
for rw in range(1, RD.max_row + 1):
    v_ = RD.cell(rw, 1).value
    if isinstance(v_, str) and rw == 1: RD.cell(rw, 1, v_.replace("v0.1.1", "v0.2.1"))
    elif isinstance(v_, str) and v_.startswith("估值："): RD.cell(rw, 1, V2["readme_val"] + "（v0.1.1 原估值說明，保留作參考：" + v_[3:] + "）")
    elif isinstance(v_, str) and v_.startswith("版本："): RD.cell(rw, 1, "版本：v0.2.1（2026-10-25：補上檔空間級距、投資建議、重點觀察；PL 頁最後新增三情境自由現金流列（淨利＋折舊 − 資本支出，不含營運資金變動）、In 頁最後新增折舊 ÷ 營收、資本支出 ÷ 營收（數值不變））。v0.2（2026-10-25：目標價改為 DCF 與 Comparable 加權、三情境延伸到 FY2032；FY2025–FY2029 數值不變）。前版：" + v_[3:])
rn = RD.max_row + 1
for t_ in V2["readme_add"]:
    RD.cell(rn, 1, t_); rn += 1
if os.environ.get("V2_ROWMAP"):
    import json as _json
    _json.dump({"XT": XT, "XH": XH, "XR": XR, "T": T, "EPSR": EPSR, "CFR": CFR, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "RECV2": RECV2},
               open(os.environ["V2_ROWMAP"], "w"), ensure_ascii=False, indent=1)

# ================= v0.2.1（2026-10-25）：自由現金流（算力收支 2026-10-25 請求；投資地圖依列標籤「自由現金流」讀取） =================
# In 頁最後新增折舊 ÷ 營收、資本支出 ÷ 營收（FY2025 實績為錨，各年沿用）；PL 頁最後新增三情境自由現金流列。不位移既有儲存格。
ir = IN.max_row + 2
IN.cell(ir, 1, "K. 自由現金流參數（v0.2.1 新增；以 FY2025 年報實績為錨，FY2026 起各年沿用；只供 PL 頁自由現金流列，不進 EPS 與目標價）").font = BOLD; ir += 1
for key_, item_, val_, s1_ in (("da_r", "折舊與攤銷 ÷ 營收", "=12.081/a25_rev", "FY2025 折舊與攤銷 120.81 億元 ÷ FY2025 營收（Q 頁 B4）；stockanalysis（S&P Global）現金流量表，2026-10-25 擷取"),
                               ("cx_r", "資本支出 ÷ 營收", "=14.757/a25_rev", "FY2025 資本支出（購置不動產、廠房及設備）147.57 億元 ÷ FY2025 營收（Q 頁 B4）；stockanalysis（S&P Global）現金流量表，2026-10-25 擷取")):
    IN.cell(ir, 1, key_); IN.cell(ir, 2, item_); IN.cell(ir, 3, "%")
    IN.cell(ir, 5, val_).number_format = "0.00%"
    IN.cell(ir, 7, "Verified（FY2025 實績）／Assumed（各年沿用）"); IN.cell(ir, 8, s1_); IN.cell(ir, 9, "FY2026 起各年＝營收 × 本比率（Assumed：資本密集度不變）")
    wb.defined_names[key_] = DefinedName(key_, attr_text=f"In!$E${ir}"); ir += 1
fr = PL.max_row + 3
PL.cell(fr, 1, "自由現金流（v0.2.1 新增；算力收支 2026-10-25 請求）：淨利（歸屬母公司，延伸表）＋折舊與攤銷（營收 × In 頁 da_r）− 資本支出（營收 × In 頁 cx_r）；不含營運資金變動。FY2030 起營收、淨利為收斂法").font = BOLD; fr += 1
hdr(PL, fr, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]); FCF_H = fr; fr += 1
FCFR = {}
for s in SC:
    PL.cell(fr, 1, f"自由現金流（淨利＋折舊 − 資本支出；不含營運資金變動）｜{SCN_T[s]}"); PL.cell(fr, 2, V2["unit"])
    for y in XY[1:]:
        c = XC[y]
        PL[f"{c}{fr}"] = f"={c}{XR[s]['ni']}+{c}{XR[s]['rev']}*da_r-{c}{XR[s]['rev']}*cx_r"; PL[f"{c}{fr}"].number_format = "#,##0.0"
        if y in FADE: PL[f"{c}{fr}"].fill = FADEFILL
    PL.cell(fr, 11, "淨利、營收引用本頁延伸表同情境列；折舊、資本支出以 FY2025 實績比率 × 營收")
    FCFR[s] = fr; fr += 1
so = SO.max_row + 1
SO.cell(so, 1, "[Verified] 廣達 FY2025 年報：折舊與攤銷 120.81 億元、資本支出 147.57 億元、營業現金流 −103.65 億元（資料商轉載合併現金流量表）"); SO.cell(so, 2, "https://stockanalysis.com/quote/tpe/2382/financials/cash-flow-statement/")
SO.cell(so + 1, 1, "[Assumed] FY2026 起折舊 ÷ 營收、資本支出 ÷ 營收沿用 FY2025 實績比率（In 頁 da_r、cx_r）；自由現金流不含營運資金變動"); SO.cell(so + 1, 2, "本模型假設")

# ================= v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（工作單；Andy 2026-10-25 14:00、14:17） =================
# 全部放在 Val、README 現有最後一列之後，不位移既有儲存格
OBS = [
       "目標價 NT$566.6（DCF 644.3、Comparable 488.9 各半），比現價 322（2026-10-22）高 76.0%，級距大於 50%；Comparable 單獨為 51.8%。主因是 EPS 不是倍數：本模型 12 個月後前瞻 EPS 43.00 比共識前瞻 EPS 28.56 高 50.6%（兩者期間差約一年；FY2027 EPS 43.61 也比法人 10-08 的 32.09 高 35.9%），同業（緯創、緯穎、鴻海、英業達、Super Micro）前瞻 P/E 中位數 11.4 倍與現價隱含 11.3 倍幾乎相同；DCF（ROE 27.2%、再投資率 14.7%）比 Comparable 高 31.8%，再把目標價往上拉。",
       "EPS 高於法人來自營收與淨利係數：FY2027 營收 5.71 兆元比法人 5.30 兆元高 7.7%，淨利 ÷ 營業利益約 0.87；改用券商隱含的 0.75 時 FY2027 EPS 降 13.3% 到 37.81，是 FY2027 最敏感的變數。3Q26 法說（11 月中）的業外與稅率是第一個檢驗點。",
       "情境分歧從 FY2028 起：FY2028 EPS A／B／C＝51.92／41.96／31.56（B、C 比 A 低 19.2%、39.2%）；份額相對值 0.75（鴻海、緯創搶單）時 FY2028 EPS −18.2%。追蹤三家 ODM 的 Vera Rubin 櫃數份額（本模型 FY2026 有效份額 13.4%）與 2027 年客供比例。"
      ]   # 重點觀察（與 README.md 相同文字；數字取自重算後的 xlsx）
_U = f"B{T['upw']}"
GR = V.max_row + 3
V.cell(GR - 1, 1, "0.4 上檔空間級距與投資建議（v0.2.1 新增；放在本頁最後以免位移既有儲存格；依「上檔空間：加權目標價」第 " + str(T["upw"]) + " 列）").font = BOLD
tset(GR, "grade", "上檔空間級距", [f'=IF({_U}>0.5,"大於 50%",IF({_U}>0.25,"25%–50%",IF({_U}>=0,"0%–25%",IF({_U}>=-0.25,"0%～−25%",IF({_U}>=-0.5,"−25%～−50%","小於 −50%")))))', None, None], "General",
     "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）；負值區間同樣重要")
tset(GR + 1, "rec", "投資建議", [f'=IF({_U}>0.5,"強力買進",IF({_U}>0.25,"買進",IF({_U}>=-0.25,"中立",IF({_U}>=-0.5,"放空","強力放空"))))', None, None], "General",
     "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）")
rn = RD.max_row + 1
RD.cell(rn, 1, f'="投資建議："&Val!B{GR + 1}&"（依上檔空間固定規則）"').font = BOLD
OBS_ROW = rn + 1
for i_, t_ in enumerate(OBS, 1):
    RD.cell(rn + i_, 1, f"重點觀察 {i_}：" + t_)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("QUANTA_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK_ROW, "SENSN": SENS_N, "IMP": IMP, "VAL": VAL, "SB": SB},
              open(os.environ["QUANTA_ROWMAP"], "w"), ensure_ascii=False, indent=1)
