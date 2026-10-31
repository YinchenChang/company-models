# 台燿（6274）個股模型 v0.1（AI 半導體 Project：載板與 PCB 次產業；高速 CCL 銅箔基板）
# 用法：python3 tuc/build_tuc.py raw.xlsx ；再以 LibreOffice headless 重算（見 README）
# 會計年度＝日曆年，FY(t)＝CY(t)。金額：十億新台幣（NT$B）；CCL 市場以美元讀 Siliconomics、以匯率 fx 換算；EPS：新台幣。
# 樣板：emc/（台光電，同屬高速 CCL）。差異：(1) 行業變數（部署 GW、每 GW CCL 金額、匯率）一律讀 Siliconomics 快照；
# (2) 敏感度改為「一次改一個 In 頁低／高值、整本重算」的加權目標價（tuc/sens_tuc.py 產生 tuc/sens.json，建置時寫入數值）；
# (3) 份額與台光電一致性：讀台光電成品 Units 頁的有效份額與 AI 營收，同一市場兩家合計不得超過 100%（建置時檢查）。
# 環境變數 TUC_OVERRIDE="鍵=值;鍵=值"：只供 sens_tuc.py 使用（把 In 頁基準欄改成低／高值）；CI 與正式建置不設。
import sys, os, glob, json
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # noqa: E402  Siliconomics 快照（行業變數單一來源）
sys.path.insert(0, os.path.join(ROOT, "tools"))
import relative_multiple as RMX   # noqa: E402
sys.path.insert(0, os.path.join(ROOT, "tools", "in_layout"))
import apply_in_layout   # noqa: E402

VER, DATE = "v0.1", "2026-10-31"
BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
SC = ("A", "B", "C")
OVR = {}
for kv in filter(None, os.environ.get("TUC_OVERRIDE", "").split(";")):
    k, v = kv.split("="); OVR[k] = float(v)
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
rel = lambda f: os.path.relpath(f, ROOT)
SNAP = f"Siliconomics 快照 {SIL.VERSION}（{SIL.FILE}）"

# ---------------- 上游：Siliconomics（行業）＋台光電成品（份額一致性檢查，唯讀） ----------------
N = SIL.NAMES
GW = {s: N[f"SIF_DeployGW_{s}"] for s in SC}          # 2026–2029
PGW = N["SIF_CCLperGW"]; MKT = {s: N[f"SIF_CCLMkt_{s}"] for s in SC}
assert all(GW[s][:2] == GW["A"][:2] for s in SC), "2026–27 部署 GW 三情境應相同"
EMC_F = glob.glob(os.path.join(ROOT, "emc", "dist", "*.xlsx")); assert len(EMC_F) == 1, EMC_F; EMC_F = EMC_F[0]
EMC = {}
for r in openpyxl.load_workbook(EMC_F, data_only=True)["Units"].iter_rows(min_row=4):
    if r[0].value in ("台光電有效份額", "台光電 AI 相關營收（實際＝MIN 需求面、產能）", "AI 高速 CCL 市場（含價格指數）"):
        EMC[r[0].value] = ([r[2 + j].value for j in range(4)], f"{rel(EMC_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert len(EMC) == 3, EMC.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E）。金額：十億新台幣（NT$B），CCL 市場與每 GW 金額為美元；產能：百萬張／月；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = ("標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。公司原文＝台燿財報（公開資訊觀測站，經 HiStock、nstock 轉載）、月營收公告與法說（經富果轉述）；"
            "券商估計皆為 Interested-party。行業變數（部署 GW、每 GW CCL 金額、匯率）讀 " + SNAP + "；本頁不另訂行業路徑。")
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
FG0730 = "https://blog.fugle.tw/post/earnings-call-6274-2026-07-30"
FG0312 = "https://blog.fugle.tw/post/earnings-call-6274-2026-03-12"
HS_IS = "https://histock.tw/stock/6274/損益表"
HS_EPS = "https://histock.tw/stock/6274/每股盈餘"
HS_REV = "https://histock.tw/stock/6274/每月營收"
NS = "https://www.nstock.tw/6274?status=6"
UA = "https://uanalyze.com.tw/articles/1773952452"
MW = "https://www.moneyweekly.com.tw/_Article?AID=247880"
UDN = "https://money.udn.com/money/story/5607/9638268"
WR = "https://wantrich.chinatimes.com/news/20260630900220-420101"
CNY = "https://news.cnyes.com/news/id/6426189"
SA = "https://stockanalysis.com/quote/tpex/6274/"

sec("A. AI 部署與每 GW 高速 CCL 金額（讀 Siliconomics 快照；情境 A／B／C 讀同字母上游）")
inp("gw_26", "部署 GW 2026（IT 電力；三情境相同）", "GW", None, GW["A"][0], None, "Derived", f"{SNAP}｜SIF_DeployGW_A[2026]", "AI IT 資本支出 ÷ 每 GW IT 資本支出（Siliconomics 推導；Tokenomics 量經 SIF_TK_*）")
inp("gw_27", "部署 GW 2027（三情境相同）", "GW", None, GW["A"][1], None, "Derived", f"{SNAP}｜SIF_DeployGW_A[2027]", "同上；顆數情境係數從 CY2028 開始")
for y, j in ((28, 2), (29, 3)):
    for s in SC:
        inp(f"gw{y}_{s}", f"情境 {s}：部署 GW 20{y}", "GW", None, GW[s][j], None, "Derived", f"{SNAP}｜SIF_DeployGW_{s}[20{y}]",
            "A＝顆數基準；B＝CY2028 ×0.8、CY2029 ×0.85；C＝CY2028–29 ×0.6（Siliconomics SIF_UnitsCoef_*；與台光電、台積電模型同一套）")
for y, j in ((26, 0), (27, 1), (28, 2), (29, 3)):
    inp(f"pgw_{y}", f"每 GW 高速 CCL 金額 20{y}（加速器／運算板＋交換器板＋UBB／中板，不含價格指數）", "US$M/GW", None, PGW[j], None, "Derived", f"{SNAP}｜SIF_CCLperGW[20{y}]",
        "2026 年三項 190／90／70（Assumed）× 世代指數（Analogy／Assumed）；元件與指數見 Siliconomics In 頁 M16（原在台光電模型）")
inp("fx", "匯率（新台幣 ÷ 美元）", "NT$/US$", *SIL.vals("fx_twd"), "Derived", f"{SNAP}｜SIF_fx_twd（與台光電、台達電模型相同）", "低／高只用於 Sens")

sec("B. 台燿份額與價格（公司變數；份額相對台光電一致：同一市場兩家合計不得超過 100%，見 Units 頁）")
for y, v in ((27, (0.95, 1.10, 1.25)), (28, (0.95, 1.15, 1.35)), (29, (0.90, 1.15, 1.40))):
    inp(f"rel_{y}", f"台燿 AI 高速 CCL 有效份額 20{y} ÷ 2026（相對值）", "x", *v, "Assumed",
        "泰國廠 2026-07 起量產、4Q26 滿載（公司，富果 2026-07-30）；常熟、泰國各 +60 萬張／月 2027Q3 開出（富果 2026-03-12）；AI ASIC 專案 2026 年貢獻 5–6%（日系外資，旺得富 2026-06-30）" if y == 27 else "同上",
        "Siliconomics 無 CCL 份額表（ShareTbl 無台燿、台光電）：本列放公司 In 頁，待併入 Siliconomics；台光電模型同期相對值 0.95／0.90／0.85（份額流失給台燿等追趕者）" if y == 27 else "待併入 Siliconomics")
for y, v in ((27, (0.97, 1.05, 1.12)), (28, (0.90, 1.00, 1.08)), (29, (0.85, 0.95, 1.03))):
    inp(f"pp_{y}", f"台燿 CCL 年均售價指數 20{y}（FY2026 年均＝1；含 2026 年中漲價的全年效果）", "x", *v, "Assumed",
        "毛利率 1Q26 25.2% → 2Q26 29.7% → 3Q26 36.9%（財報），2026 年中漲價使 FY2026 年均價低於年底水準；美系外資預期 M7 以上 CCL 第二波季漲 10–15%（經濟日報 2026-07-20）" if y == 27 else "同上",
        "高階 CCL 價格屬行業變數，Siliconomics 尚無價格路徑：待併入 Siliconomics；台光電模型 2027–29 用 1.00／0.97／0.93（以 FY2026 年均為基期，台光電漲價較早）" if y == 27 else "新產能到位後回落（Assumed）；待併入 Siliconomics")

sec("C. 產能（全產全銷：營收上限＝平均月產能 × 12 × 每張營收）")
inp("cap_26", "平均月產能 2026", "百萬張／月", 2.25, 2.35, 2.45, "Derived", "集團 CCL 月產能逾 220 萬張（官方資料，理財周刊 2026-09-10）；泰國廠 2026-07 起量產、4Q 滿載（富果 2026-07-30）",
    "市場估 2026 年底約 260 萬張（理財周刊，非公司指引）；年均＝220 萬 → 260 萬的時間加權")
inp("cap_27", "平均月產能 2027", "百萬張／月", 2.75, 3.10, 3.30, "Derived", "常熟、泰國各 +60 萬張／月，2027Q3 投產（董事會 2026-03-11，富果 2026-03-12）；市場估 2027 年底趨近 380 萬（理財周刊 2026-09-10）",
    "260 萬 ＋ 120 萬 × 5／12；低值反映設備交期延遲（泰國一期曾延兩個月）")
inp("cap_28", "平均月產能 2028", "百萬張／月", 3.80, 4.30, 4.70, "Derived", "2028–29 年新增 195 萬張／月：中山 +120 萬、泰國二期 +75 萬（董事會，富果 2026-07-30）；可轉債 100 億元籌資",
    "380 萬 ＋ 195 萬於 2028–29 線性開出的年均（Assumed 時程）")
inp("cap_29", "平均月產能 2029", "百萬張／月", 4.30, 5.30, 5.75, "Assumed", "同上；全部開出為 575 萬張／月", "低值＝新廠延後一年")
for y, v in ((27, (1.08, 1.20, 1.35)), (28, (1.12, 1.32, 1.55)), (29, (1.15, 1.42, 1.70))):
    inp(f"mix_{y}", f"每張營收組合指數 20{y}（M7／M8 → M8+、層數；不含價格；2026＝1）", "x", *v, "Assumed",
        "ELL／SLL（M7、M8）占營收：2021 年 10% → 4Q25 36% → 1Q26 38% → 2Q26 39%（公司，富果 2026-03-12、07-30）" if y == 27 else "同上",
        "AI ASIC 專案用 Low Dk2 玻纖布、HVLP4 銅箔（理財周刊 2026-09-10）；台光電模型組合指數 1.25／1.40／1.50（M9 較早量產）" if y == 27 else "—")

sec("D. 營收結構錨點（公司揭露的損耗等級組合 × 營收）")
inp("ai_25", "2025 年高速材料（ELL／SLL，M7 以上）÷ 營收", "%", 0.28, 0.32, 0.36, "Assumed", "公司：ELL／SLL 4Q25 36%（富果 2026-03-12）；全年比例未揭露，前三季低於 4Q（Assumed）", "只用於 FY2025 欄的子列拆分")
inp("ai_26", "2026 年高速材料（ELL／SLL，M7 以上）÷ 營收（＝AI 相關的代理）", "%", 0.38, 0.42, 0.48, "Interested-party", "公司：1Q26 38%、2Q26 39%（富果 2026-07-30）；3Q26 毛利率跳升至 36.9%，高階占比續升（Assumed）",
    "讀法：AI 伺服器、800G 交換器、AI ASIC 用 M7 以上；M4／M6 的一般伺服器與 400G 以下網通列入其他等級（Assumed 對應）")
for y, v in ((27, (0.0, 0.10, 0.18)), (28, (0.0, 0.07, 0.14)), (29, (0.0, 0.05, 0.10))):
    inp(f"g_nai{y}", f"其他等級營收年增 20{y}（VLL／LL、非低損耗、特殊應用）", "%", *v, "Assumed",
        "特殊應用（厚銅、電源）占營收 3Q25 8% → 2Q26 14%，2025 年年增 49%（公司，富果 2026-03-12、07-30）" if y == 27 else "同上",
        "非低損耗占比 18% → 16%（1Q26 → 2Q26）持續下降（公司）" if y == 27 else "—")

sec("E. 毛利率")
inp("d_ai", "高速材料毛利率高於其他等級的幅度", "pt", 0.05, 0.08, 0.12, "Assumed", "等級別毛利率未揭露：未找到；毛利率隨 M7 以上占比與漲價上升：FY2025 22.7% → 3Q26 36.9%（財報）", "與台光電模型同值（Analogy）")

sec("F. FY2026 錨點（1H26、3Q26 財報實績＋4Q26 假設；3Q26 財報已公布，法說日期未找到）")
inp("q4g", "4Q26 營收季增", "%", 0.0, 0.05, 0.12, "Assumed", "公司未給 4Q 指引：未找到；9 月營收 61.26 億 × 3＝183.8 億（季增 2.4%，月營收公告）", "泰國廠 4Q 滿載（公司，富果 2026-07-30）")
inp("q4_gm", "4Q26 毛利率", "%", 0.34, 0.36, 0.38, "Assumed", "3Q26 36.93%（財報；nstock 36.93%）", "銅箔、玻纖布漲價壓縮毛利（公司，富果 2026-03-12）")
inp("q4_opx", "4Q26 營業費用 ÷ 營收", "%", 0.055, 0.06, 0.065, "Assumed", "3Q26 1.060 ÷ 17.949＝5.9%（財報）", "2Q26 6.2%、FY2025 8.4%")
inp("q4_nop", "4Q26 業外收益", "NT$B", -0.1, 0.03, 0.15, "Assumed", "3Q26 0.026、2Q26 0.069（財報，稅前 − 營業利益）", "—")
inp("q_nir", "4Q26 稅後淨利 ÷ 稅前", "%", 0.66, 0.69, 0.75, "Derived", "3Q26 3.857 ÷ 5.594＝0.689；2Q26 0.683（財報）", "FY2025 0.754（財報）")
inp("q_sh", "股數（十億股；EPS 換算）", "十億股", None, "=3.856631/12.88", None, "Derived", "3Q26 稅後淨利 38.57 億 ÷ EPS 12.88（HiStock、nstock）", "stockanalysis：流通股數 288.78 百萬股；可轉債轉換使加權股數上升")

sec("G. 營業費用、業外、稅與股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.45, 0.6, 1.0, "Derived", "FY2026 錨點實測約 0.46（Q 頁 B26）", "與台光電模型同值")
inp("nop_f", "業外收益（每年）", "NT$B", -0.3, 0.15, 0.5, "Assumed", "FY2025 0.175、1H26 0.115（財報）", "—")
inp("nir_f", "稅後淨利 ÷ 稅前", "%", 0.66, 0.70, 0.76, "Derived", "1H26 0.680、3Q26 0.689、FY2025 0.754（財報）", "—")
inp("sh_g", "股數年變動", "%", 0.0, 0.01, 0.025, "Assumed", "董事會通過可轉債面額上限 100 億元、轉換價不低於 110%（優分析 2026-07-30）：全數轉換約稀釋 2.5%", "—")
inp("mf", "AI 高速 CCL 市場係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "Siliconomics 市場 × 本係數", "±20%")

sec("H. 共識與券商（只作 Recon 對帳，不作校準）")
inp("c_eps26f", "FactSet 共識 FY2026 EPS 中位數（2026-08-04，15 位）", "NT$", None, 36.05, None, "Interested-party", MW + "（理財周刊轉引 FactSet）", "營收中位數約 588.9 億")
inp("c_eps27f", "FactSet 共識 FY2027 EPS 中位數（2026-08-04）", "NT$", None, 70.47, None, "Interested-party", MW, "低 53.91、高 100.07；營收中位數約 1,007 億")
inp("c_rev26f", "FactSet 共識 FY2026 營收中位數（2026-08-04）", "NT$B", None, 58.89, None, "Interested-party", MW, "—")
inp("c_rev27f", "FactSet 共識 FY2027 營收中位數（2026-08-04）", "NT$B", None, 100.7, None, "Interested-party", MW, "—")
inp("c_eps28f", "FactSet 共識 FY2028 EPS（2026-04-20，單一預估）", "NT$", None, 81.76, None, "Interested-party", CNY, "期間早、只有一家，僅供量級參考")
inp("b_eps26u", "美系券商 FY2026 EPS（2026-07-20）", "NT$", None, 38.8, None, "Interested-party", UDN + "（經濟日報，未具名美系券商）", "3Q26 毛利率估 32.3%")
inp("b_eps27u", "美系券商 FY2027 EPS（2026-07-20）", "NT$", None, 93.5, None, "Interested-party", UDN, "—")
inp("b_eps28u", "美系券商 FY2028 EPS（2026-07-20）", "NT$", None, 168.7, None, "Interested-party", UDN, "—")

sec("I. 價格")
inp("px", "股價（2026-10-30 收盤）", "NT$", None, 1265, None, "Interested-party", "HiStock 6274（10-30 13:30 收盤 1,265，−8.99%）", "nstock：10-30 收盤 1,265（前收 1,390）")

sec("J. 目標價（DCF 與 Comparable 加權）與收斂法參數（黃底＝Assumed，Andy 可改）")
for s_, w_, nm_ in (("A", 0.3, "順利"), ("B", 0.5, "平平"), ("C", 0.2, "不如預期")):
    inp(f"tp_w{s_}", f"目標價情境權重：情境 {s_}（{nm_}）", "權重", None, w_, None, "Assumed", "Andy 可改（2026-10-25 目標價規格預設 A／B／C＝0.3／0.5／0.2）", "三個權重合計應為 1")
inp("tp_r", "折現率／要求報酬（DCF）", "%", 0.09, 0.10, 0.11, "Assumed", "Andy 可改（工作單預設 10%；美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（DCF 永續；情境 B 的 FY2032 營收年增收斂目標）", "%", 0.03, 0.04, 0.05, "Assumed", "Andy 可改（工作單預設 4%）", "Sens 頁 ±1pt（同時移動 PL 收斂目標 fade_g）")
inp("tp_roe", "ROE（近三年平均）", "%", None, "=AVERAGE(0.0703,0.2010,0.2069)", None, "Derived", "stockanalysis 比率頁（Return on Equity）：FY2023 7.03%、FY2024 20.10%、FY2025 20.69%", "單一來源；近 12 月 30.74%（同站，2026-10-26）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", 0.3, 0.5, 0.7, "Assumed", "Andy 可改（2026-10-25 11:53：DCF／Comparable 各半）", "Sens 頁 0.3／0.7")
inp("fade_gA", "情境 A（順利）：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂，FY2032 達到")
inp("fade_gB", "情境 B（平平）：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C（不如預期）：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("nim10_f", "十年平均淨利率的替代：最後一年（FY2029）淨利率 ×（本值）", "x", None, 0.8, None, "Assumed", "十年（2016–2025）平均淨利率：未找到（2016–2017 年財報數字未取得）；依工作單以最後一年 × 0.8 替代", "淨利率收斂目標：A＝FY2029、B＝中點、C＝替代值")
inp("cons_tp", "市場共識目標價（賣方平均）", "NT$", None, 2213, None, "Interested-party", "stockanalysis 個股頁 Analyst price target 2,213（頁面時間 2026-10-22）", "FactSet 2026-04-20 目標價 900（較舊）；日系外資 2,115（2026-06-30）")
inp("cons_feps", "共識前瞻 12 月 EPS（＝價格 ÷ 前瞻 P/E）", "NT$", None, "=1420/26.37", None, "Interested-party", "stockanalysis 個股頁（2026-10-22）：價格 1,420 ÷ Forward PE 26.37", "賣方共識，只作對帳與差距拆解，不作校準")

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    if key in OVR: base = OVR[key]
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}"); INROW[key] = r
    if key.startswith(("tp_", "fade_g", "nim10")): IN.cell(r, 5).fill = KEY
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
INLIST = [x for x in rows if x[0] != "sec"]

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣；股數為十億股）。藍字＝輸入；黑字＝公式。實績取自台燿合併財報（公開資訊觀測站，經 HiStock 損益表、nstock 轉載）與月營收公告"
Q["A2"] = "FY2026 錨點＝1H26 財報＋3Q26 財報（10 月底公布）＋4Q26 假設（營收＝3Q × (1＋季增)、毛利率、營業費用率、業外、稅後 ÷ 稅前）。10 月營收（約 11-10）與 3Q26 法說（日期未找到）後更新"
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（實績）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1", "來源2／說明"])
S1 = "HiStock 損益表（季，千元；公開資訊觀測站轉載）：" + HS_IS
S2 = "nstock 3Q26：毛利率 36.93%、營益率 31.02%、EPS 12.88；1H26：優分析 2026-07-30（營收 243.56 億、EPS 12.4）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "營收", {"B": "=6.372144+6.780435+8.062873+9.124783", "C": "=10.053966+14.301042", "D": 17.948807, "E": "=D4*(1+q4g)"}, tag="Verified／Assumed（4Q 季增）",
     s2="月營收：FY2025 303.44 億、1–9 月 423.04 億（HiStock 月營收 " + HS_REV + "）；財報季營收合計 303.40 億")
qrow(5, "毛利", {"B": "=1.533022+1.435244+1.914793+2.014734", "C": "=2.528558+4.249571", "D": 6.628111, "E": "=E4*q4_gm"}, tag="Verified／Assumed",
     s2="1H26 毛利 67.78 億、毛利率 27.83%（工商時報 2026-07-31）；stockanalysis FY2025 毛利 68.98 億")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="3Q26 36.93%（nstock）")
qrow(7, "營業費用", {"B": "=B5-B8", "C": "=C5-C8", "D": "=D5-D8", "E": "=E4*q4_opx"}, tag="Derived（毛利 − 營業利益）／Assumed")
qrow(8, "營業利益", {"B": "=0.935418+0.852572+1.265539+1.290130", "C": "=1.827536+3.357936", "D": 5.568008, "E": "=E5-E7"}, tag="Verified", s2="1H26 營業利益 51.85 億、營益率 21.29%（工商時報 2026-07-31）")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="3Q26 31.02%（nstock）")
qrow(10, "業外收益", {"B": "=B11-B8", "C": "=C11-C8", "D": "=D11-D8", "E": "=q4_nop"}, tag="Derived（稅前 − 營業利益）／Assumed")
qrow(11, "稅前利益", {"B": "=0.934823+0.839586+1.317663+1.426540", "C": "=1.873851+3.426937", "D": 5.594157, "E": "=E8+E10"}, tag="Verified")
qrow(12, "稅後淨利（歸屬母公司）", {"B": "=0.671947+0.651731+1.002514+1.082957", "C": 3.602, "D": 3.856631, "E": "=E11*q_nir"}, tag="Verified／Derived（4Q）",
     s2="1H26 歸屬母公司淨利 36.02 億（理財周刊 2026-09-10）；HiStock 季稅後淨利 1Q 12.60、2Q 23.42 億")
qrow(13, "稅後 ÷ 稅前", {c: f"={c}12/{c}11" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(14, "股數（十億股）", {"B": "=B12/12.13", "C": "=C12/12.4", "D": "=q_sh", "E": "=q_sh"}, ann="=q_sh", tag="Derived", s2="淨利 ÷ EPS（加權平均；可轉債轉換使股數上升）")
qrow(15, "EPS", {c: f"={c}12/{c}14" for c in "BCDE"}, fmt="0.00", ann="=F12/F14", tag="Derived", s1="淨利 ÷ 股數", s2="FY2026 以 3Q26 股數換算（低估 1H 股數較少的效果，差約 1%）")
qrow(16, "  公司揭露 EPS", {"B": 12.13, "C": 12.4, "D": 12.88}, fmt="0.00", ann="", tag="Verified", s1="HiStock 每股盈餘 " + HS_EPS, s2="1Q26 4.36、2Q26 8.02、3Q26 12.88；FY2025 季 2.43／2.36／3.58／3.76（nstock 同值）")
qrow(17, "高速材料營收（ELL／SLL，M7 以上；＝營收 × 占比）", {"B": "=B4*ai_25"}, ann="=F4*ai_26", tag="Derived／Interested-party", s1="公司損耗等級組合（In 頁 D 節）", s2="各季金額未揭露：未找到")
qrow(18, "其他等級營收（＝營收 − 高速材料）", {"B": "=B4-B17"}, ann="=F4-F17", tag="Derived", s1="—", s2="VLL／LL（M4、M6）、非低損耗、特殊應用（厚銅、電源）")
Q.cell(23, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(24, 1, "FY2026 營收年增"); Q["B24"] = "=F4/B4-1"
Q.cell(25, 1, "FY2026 營業費用年增"); Q["B25"] = "=F7/B7-1"
Q.cell(26, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B26"] = "=B25/B24"
Q.cell(27, 1, "FY2026 營收 ÷ FactSet 共識（08-04）− 1"); Q["B27"] = "=F4/c_rev26f-1"
Q.cell(28, 1, "FY2026 EPS ÷ 美系券商（07-20）− 1"); Q["B28"] = "=F15/b_eps26u-1"
Q.cell(29, 1, "3Q26 營收季增（實績）"); Q["B29"] = "=D4/14.301042-1"
Q.cell(30, 1, "季報合計 − 公司揭露 EPS（FY2025；四捨五入差）"); Q["B30"] = "=B15-B16"
for rr in range(24, 31): Q[f"B{rr}"].number_format = "0.0%" if rr not in (26, 30) else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 12), ("sh", 14), ("eps", 15), ("ai", 17), ("nai", 18)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$F${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 90; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
gwk = lambda y, s: {2026: "gw_26", 2027: "gw_27", 2028: f"gw28_{s}", 2029: f"gw29_{s}"}[y]
def engine(ws, r0, title, s):
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 高速 CCL 市場（Siliconomics）與台燿份額（需求面；D 欄＝FY2026 校準年）")
    row("gw", "部署 GW（IT 電力）", "GW", {c_(y): f"={gwk(y, s)}" for y in BU}, fmt="0.0", note=f"Siliconomics SIF_DeployGW_{s}")
    row("pgw", "每 GW 高速 CCL 金額（不含價格指數）", "US$M/GW", {c_(y): f"=pgw_{y % 100}" for y in BU}, fmt="#,##0", note="Siliconomics SIF_CCLperGW")
    row("mku", "AI 高速 CCL 市場（美元；不含價格指數）", "US$B", {c_(y): f"={c_(y)}{R['gw']}*{c_(y)}{R['pgw']}/1000" + ("" if y == 2026 else "*mf") for y in BU},
        note=f"＝ Siliconomics SIF_CCLMkt_{s}（Recon 核對）；FY2027 起 × mf（Sens 用，基準 1）")
    row("mkn", "AI 高速 CCL 市場（新台幣）", "NT$B", {c_(y): f"={c_(y)}{R['mku']}*fx" for y in BU}, fmt="#,##0.0")
    row("pp", "台燿年均售價指數（FY2026＝1）", "x", {"D": "=1", **{c_(y): f"=pp_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh0", "台燿有效份額（FY2026 校準）＝ 高速材料營收 ÷ 市場", "%", {"D": f"=a26_ai/D{R['mkn']}"}, fmt="0.0%",
        note="Derived：分子為公司 ELL／SLL 占比 × 營收；分母為 AI 高速 CCL 市場（含 ASIC）")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"=rel_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "台燿有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("aid", "高速材料營收（需求面）＝ 市場 × 份額 × 價格指數", "NT$B", {c_(y): f"={c_(y)}{R['mkn']}*{c_(y)}{R['sh']}*{c_(y)}{R['pp']}" for y in BU}, fmt="#,##0.0")
    sec_("B. 產能上限（全產全銷）與營收")
    row("cap", "平均月產能", "百萬張／月", {c_(y): f"=cap_{y % 100}" for y in BU}, fmt="0.00")
    row("rps0", "每張營收（FY2026 校準）＝ 錨點營收 ÷（月產能 × 12）", "NT$／張", {"D": f"=a26_rev/(D{R['cap']}*12)*1000"}, fmt="#,##0",
        note="Derived：全產全銷，FY2026 營收全部由產能吸收；含膠片與所有等級")
    row("mix", "每張營收組合指數（FY2026＝1）", "x", {"D": "=1", **{c_(y): f"=mix_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("rcap", "產能上限營收 ＝ 月產能 × 12 × 每張營收 × 組合指數 × 價格指數", "NT$B",
        {c_(y): f"={c_(y)}{R['cap']}*12*$D${R['rps0']}*{c_(y)}{R['mix']}*{c_(y)}{R['pp']}/1000" for y in BU}, fmt="#,##0.0")
    row("nai", "其他等級營收（VLL／LL、非低損耗、特殊應用）", "NT$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+g_nai{y % 100})" for y in BU[1:]}}, fmt="#,##0.0")
    row("ai", "高速材料營收（ELL／SLL，M7 以上；AI 相關）＝ MIN（需求面、產能上限 − 其他等級）", "NT$B", {"C": "=a25_ai", "D": f"=D{R['aid']}",
        **{c_(y): f"=MIN({c_(y)}{R['aid']},{c_(y)}{R['rcap']}-{c_(y)}{R['nai']})" for y in BU[1:]}}, fmt="#,##0.0", note="產能不足時，其他等級照成長率出貨、高速材料取剩餘產能（Assumed）")
    row("bind", "產能是否受限", "", {c_(y): f'=IF({c_(y)}{R["aid"]}>{c_(y)}{R["rcap"]}-{c_(y)}{R["nai"]}+0.0001,"受限","未受限")' for y in BU[1:]}, fmt="General")
    row("rev", "營收（單一應報導部門：銅箔基板與膠片）", "NT$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.0", note="官方分部（segments.json）；子列＝高速材料＋其他等級")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("segchk", "檢查：分部合計 − 合併營收", "NT$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}-{c_(y)}{R['rev']}" for y in FYS}, fmt="0.000",
        note="子列（高速材料＋其他等級）合計 − 單一部門營收；FY2025、FY2026 單一部門營收＝財報（Q 頁）")
    row("aish", "高速材料 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("util", "需求面 ÷ 產能上限（高速材料需求面＋其他等級）", "%", {c_(y): f"=({c_(y)}{R['aid']}+{c_(y)}{R['nai']})/{c_(y)}{R['rcap']}" for y in BU}, fmt="0.0%",
        note="大於 100%＝需求超過產能（營收由產能決定）")
    row("pgwe", "每 GW 台燿高速材料營收", "US$M/GW", {c_(y): f"={c_(y)}{R['ai']}/fx/{c_(y)}{R['gw']}*1000" for y in BU}, fmt="#,##0")
    sec_("C. 損益")
    row("gmn", "其他等級毛利率 ＝ 錨點毛利率 − 差距 × 高速材料占比（FY2026 校準）", "%", {"D": f"=a26_gp/D{R['rev']}-d_ai*D{R['aish']}"}, fmt="0.0%",
        note="Derived：讓 FY2026 毛利率等於錨點；高速材料毛利率＝其他＋差距")
    row("gma", "高速材料毛利率 ＝ 1 −（1 − 校準值）÷ 價格指數", "%", {"D": f"=D{R['gmn']}+d_ai", **{c_(y): f"=1-(1-$D${R['gmn'] + 1})/{c_(y)}{R['pp']}" for y in BU[1:]}}, fmt="0.0%",
        note="Derived：每張成本不變，價格升降全部落在毛利（組合升級的毛利效果未另計）")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"={c_(y)}{R['ai']}*{c_(y)}{R['gma']}+{c_(y)}{R['nai']}*$D${R['gmn']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("opx", "營業費用", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in BU[1:]}}, fmt="#,##0.0")
    row("oi", "營業利益", "NT$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.0")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("nop", "業外收益", "NT$B", {c_(y): "=nop_f" for y in BU[1:]}, fmt="#,##0.00", note="FY2025、FY2026 含在淨利錨點內")
    row("ni", "稅後淨利（歸屬母公司）", "NT$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}+{c_(y)}{R['nop']})*nir_f" for y in BU[1:]}}, fmt="#,##0.0")
    row("nim", "淨利率", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.0000")
    row("eps", "EPS", "NT$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 台燿損益（十億新台幣；EPS 新台幣；歸屬母公司）。FY2025A 實績、FY2026 錨點（1H26、3Q26 財報＋4Q26 假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要；最下方延伸到 FY2032（FY2030 起為收斂法）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：順利（Siliconomics 情境 A：顆數基準）", "B": "情境 B：平平（Siliconomics 情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85）", "C": "情境 C：不如預期（Siliconomics 情境 C：CY2028–29 顆數 ×0.6）"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], s); r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  高速材料營收（M7 以上；AI 相關）", "NT$B", "#,##0.0"),
                                ("nai", "  其他等級營收", "NT$B", "#,##0.0"), ("util", "  需求面 ÷ 產能上限", "%", "0.0%"), ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"),
                                ("ni", "稅後淨利（歸屬母公司）", "NT$B", "#,##0.0"), ("eps", "EPS", "NT$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg", "util") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收、高速材料、毛利 − 錨點").font = BOLD; sr += 1
CHK = sr
PL.cell(sr, 1, "  情境 A：營收差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; sr += 1
PL.cell(sr, 1, "  情境 A：高速材料差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['ai']}-a26_ai"; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['gma']}*D{PB['A']['ai']}+D{PB['A']['gmn']}*D{PB['A']['nai']}-a26_gp"; sr += 1
PL.cell(sr, 1, "  情境 A：產能上限 − 營收（NT$B；FY2026 校準恆等）"); PL[f"D{sr}"] = f"=D{PB['A']['rcap']}-D{PB['A']['rev']}"
for rr in range(CHK, sr + 1): PL[f"D{rr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 66; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# --- PL 延伸表（三情境到 FY2032；FY2030 起收斂法） ---
XY = (2025, 2026, 2027, 2028, 2029, 2030, 2031, 2032); XC = {y: "CDEFGHIJ"[i] for i, y in enumerate(XY)}; FADE = (2030, 2031, 2032)
xr = PL.max_row + 3; XT = xr
PL.cell(xr, 1, "延伸：三情境到 FY2032。FY2025–FY2029 引用上方推導區塊；FY2030 起為收斂法（灰底），非驅動因子推導：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 線性收斂到收斂目標，FY2032 達到").font = BOLD
xr += 1; XR = {}; XH = None
for s in SC:
    E = PB[s]; R = {}
    PL.cell(xr, 1, f"延伸｜{SCN_T[s]}").font = BOLD; xr += 1
    hdr(PL, xr, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    XH = XH or xr; xr += 1
    for key, lab, unit in (("rev", "營收（延伸）", "NT$B"), ("g", "  營收年增", "%"), ("ni", "淨利（延伸）", "NT$B"), ("nim", "  淨利率", "%"),
                           ("sh", "股數（延伸）", "十億股"), ("eps", "EPS（延伸）", "NT$"), ("epsg", "  EPS 年增", "%"),
                           ("tg", "收斂目標：FY2032 營收年增", "%"), ("tm", "收斂目標：FY2032 淨利率", "%")):
        PL.cell(xr, 1, lab); PL.cell(xr, 2, unit); R[key] = xr; xr += 1
    G = lambda k: f"$G${R[k]}"
    for y in XY:
        c = XC[y]; pc = XC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={c}{E['rev']}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            PL[f"{c}{R['ni']}"] = f"={c}{E['ni']}"; PL[f"{c}{R['sh']}"] = f"={c}{E['sh_']}"; PL[f"{c}{R['eps']}"] = f"={c}{E['eps']}"
            PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"={G('g')}+($C${R['tg']}-{G('g')})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"={G('nim')}+($C${R['tm']}-{G('nim')})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['sh']}"] = f"={pc}{R['sh']}*(1+sh_g)"
            PL[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh']}"
        if y > 2025: PL[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("rev", "#,##0.0"), ("g", "0.0%"), ("ni", "#,##0.0"), ("nim", "0.0%"), ("sh", "0.0000"), ("eps", "0.00"), ("epsg", "0.0%")):
            PL[f"{c}{R[key]}"].number_format = fm
            if y in FADE: PL[f"{c}{R[key]}"].fill = FADEFILL
    PL[f"C{R['tg']}"] = f"=fade_g{s}"; PL[f"C{R['tg']}"].number_format = "0.0%"
    PL[f"C{R['tm']}"] = {"A": f"=G{R['nim']}", "B": f"=G{R['nim']}*(1+nim10_f)/2", "C": f"=G{R['nim']}*nim10_f"}[s]; PL[f"C{R['tm']}"].number_format = "0.0%"
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導（上游市場層只到 2029）")
    PL.cell(R["tg"], 11, {"A": "In 頁 fade_gA＝tp_g＋2pt", "B": "In 頁 fade_gB＝tp_g", "C": "In 頁 fade_gC＝tp_g−2pt"}[s])
    PL.cell(R["tm"], 11, {"A": "＝本情境 FY2029 淨利率", "B": "＝FY2029 淨利率與十年平均替代值（FY2029 × nim10_f）的中點", "C": "＝十年平均替代值（FY2029 × nim10_f；十年平均未找到）"}[s])
    PL.cell(R["sh"], 11, "FY2030 起股數年變動＝sh_g")
    XR[s] = R; xr += 1
PL.column_dimensions["I"].width = 14; PL.column_dimensions["J"].width = 14; PL.column_dimensions["K"].width = 60

# ---------------- Units：市場、份額、與台光電的一致性 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — AI 高速 CCL 市場（Siliconomics）、台燿份額與產能；份額相對台光電一致（同一市場兩家合計 ≤ 100%）。台光電數值讀 " + rel(EMC_F) + "（唯讀）"
A_ = PB["A"]; cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
EMC_SH = EMC["台光電有效份額"]; EMC_AI = EMC["台光電 AI 相關營收（實際＝MIN 需求面、產能）"]
for key, lab, unit, f, fmt, note in (
        ("gw", "部署 GW（情境 A）", "GW", lambda y: f"=PL!{FC[y]}{A_['gw']}", "0.0", "Siliconomics SIF_DeployGW_A"),
        ("pgw", "每 GW 高速 CCL 金額（不含價格指數）", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgw']}", "#,##0", "Siliconomics SIF_CCLperGW"),
        ("mku", "AI 高速 CCL 市場（情境 A；不含價格指數）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}", "#,##0.00", "＝ SIF_CCLMkt_A"),
        ("sh", "台燿有效份額（需求面）", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以公司 ELL／SLL 占比校準；之後 × 相對值"),
        ("aid", "台燿高速材料營收（需求面）", "US$B", lambda y: f"=PL!{FC[y]}{A_['aid']}/fx", "#,##0.00", "含台燿價格指數"),
        ("ai", "台燿高速材料營收（實際＝MIN 需求面、產能）", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx", "#,##0.00", ""),
        ("sha", "台燿實際份額（實際營收 ÷ 市場 × 價格指數）", "%", lambda y: f"=PL!{FC[y]}{A_['ai']}/PL!{FC[y]}{A_['mkn']}/PL!{FC[y]}{A_['pp']}", "0.0%", "產能受限時低於需求面份額"),
        ("pgwe", "每 GW 台燿高速材料營收", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgwe']}", "#,##0", ""),
        ("cap", "平均月產能", "百萬張／月", lambda y: f"=PL!{FC[y]}{A_['cap']}", "0.00", "公司擴產計畫"),
        ("rps", "每張營收（含組合與價格指數）", "NT$／張", lambda y: f"=PL!$D${A_['rps0']}*PL!{FC[y]}{A_['mix']}*PL!{FC[y]}{A_['pp']}", "#,##0", "FY2026 校準"),
        ("util", "需求面 ÷ 產能上限", "%", lambda y: f"=PL!{FC[y]}{A_['util']}", "0.0%", "大於 100%＝產能受限"),
        ("esh", "台光電有效份額（需求面；台光電成品）", "%", lambda y: EMC_SH[0][y - 2026], "0.0%", EMC_SH[1]),
        ("eai", "台光電 AI 相關營收（實際；台光電成品）", "US$B", lambda y: EMC_AI[0][y - 2026], "#,##0.00", EMC_AI[1]),
        ("sum", "份額一致性：台光電＋台燿 需求面份額合計", "%", lambda y: f"={cyc(y)}{{esh}}+{cyc(y)}{{sh}}", "0.0%", "同一市場兩家合計須 ≤ 100%（建置時檢查）"),
        ("suma", "份額一致性：台光電＋台燿 實際營收 ÷ 市場（不含價格指數）", "%", lambda y: f"=({cyc(y)}{{eai}}+{cyc(y)}{{ai}}/PL!{FC[y]}{A_['pp']})/{cyc(y)}{{mku}}", "0.0%", "台燿營收除以台燿價格指數；台光電營收含台光電價格指數（口徑近似）"),
        ("oth", "其他廠商（斗山、生益、松下、聯茂等）＝ 1 − 兩家需求面份額", "%", lambda y: f"=1-{cyc(y)}{{sum}}", "0.0%", "")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit); UR[key] = ur
    for y in BU:
        v = f(y)
        if isinstance(v, str): v = v.replace("{esh}", "").replace("{sh}", "").replace("{eai}", "").replace("{ai}", "").replace("{mku}", "").replace("{sum}", "")
        UN[f"{cyc(y)}{ur}"] = v; UN[f"{cyc(y)}{ur}"].number_format = fmt
        if not isinstance(v, str): UN[f"{cyc(y)}{ur}"].font = BLUE
    UN.cell(ur, 7, note); ur += 1
for y in BU:   # 以列號代入（上面先寫空位，這裡補列號）
    c = cyc(y)
    UN[f"{c}{UR['sum']}"] = f"={c}{UR['esh']}+{c}{UR['sh']}"
    UN[f"{c}{UR['suma']}"] = f"=({c}{UR['eai']}+{c}{UR['ai']}/PL!{FC[y]}{A_['pp']})/{c}{UR['mku']}"
    UN[f"{c}{UR['oth']}"] = f"=1-{c}{UR['sum']}"
ur += 1
UN.cell(ur, 1, "讀法：每 GW CCL 金額與市場讀 Siliconomics（情境 A 列示；B、C 見 PL 頁）；絕對水準只影響 FY2026 校準出的份額。Siliconomics 尚無 CCL 份額表，台燿份額相對值放本模型 In 頁（待併入 Siliconomics）；台光電份額讀台光電成品，兩家需求面份額合計 2026–29 年皆低於 100%。").font = BOLD
UNOTE = ur
UN.column_dimensions["A"].width = 66; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 70

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — 目標價（第 0 節：DCF 與 Comparable 加權）、現價隱含倍數與十年 P/E 分位檢查（新台幣）。本版未做先驗，無 Andy 合理 P/E 對照"
V["A2"] = "三情境欄 B／C／D＝A 順利／B 平平／C 不如預期（讀 Siliconomics 同字母情境）。現價＝In 頁 px（2026-10-30 收盤）"
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
t0 = 4
V.cell(t0, 1, "0. 目標價（DCF 與 Comparable 加權）").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["h1"] = t0 + 2
r_ = t0 + 3
V.cell(r_, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{r_}"] = "PL 頁延伸表；FY2030 起為收斂法"; T["eps_h"] = r_; r_ += 1
EPSR = {}
for y in YRS6:
    tset(r_, f"eps{y}", f"  FY{y}", [f"=PL!{XC[y]}{XR[s]['eps']}" for s in SC], "0.00"); EPSR[y] = r_; r_ += 1
V.cell(r_, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = r_; r_ += 1
tset(r_, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）"); RV = r_; r_ += 1
CFR = {}
for y in YRS6:
    tset(r_, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], "0.00"); CFR[y] = r_; r_ += 1
c1, c6 = CFR[2027], CFR[2032]
tset(r_, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], "#,##0",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027；低於 0 時取 0"); r_ += 1
tset(r_, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{T['dcfv']}*(1+tp_r)" for c in "BCD"], "#,##0"); r_ += 1
tset(r_, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{T['dcf12']}+tp_wB*C{T['dcf12']}+tp_wC*D{T['dcf12']}", None, None], "#,##0", "Σ tp_wA／wB／wC × 情境價值 ×(1＋tp_r)；權重見 In 頁"); r_ += 2
V.cell(r_, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; r_ += 1
hdr(V, r_, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["ph"] = r_; r_ += 1
PEERS = [("台光電（Elite Material）", "2383.TW", 32.17, "2026-09-04", "stockanalysis 個股頁 Forward PE（頁面快取 2026-09-04；較新資料未找到）"),
         ("聯茂（ITEQ）", "6213.TW", 27.49, "2026-10-22", "stockanalysis 個股頁 Forward PE（分析師共識；價格日期如左）"),
         ("金像電（Gold Circuit）", "2368.TW", 19.57, "2026-10-22", "stockanalysis 個股頁 Forward PE（分析師共識；價格日期如左）")]
P1 = r_
for nm, tk, pe, dt, src in PEERS:
    V.cell(r_, 1, nm); V.cell(r_, 2, tk); cc = V.cell(r_, 3, pe); cc.font = BLUE; cc.number_format = "0.0"
    V.cell(r_, 4, dt); V.cell(r_, 5, src); V.cell(r_, 6, "Interested-party（賣方共識）"); r_ += 1
P2 = r_ - 1; T["p1"], T["p2"] = P1, P2
tset(r_, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{P1}:C{P2})", None, None], "0.0",
     "高速 CCL 與 AI 伺服器 PCB 同業：台光電、聯茂（CCL）、金像電（下游 PCB）。生益科技（600183.SS）為 A 股，依規則不列；斗山電子材料、松下非單獨上市，未列"); r_ += 1
tset(r_, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*({2/12}*{c}{EPSR[2027]}+{10/12}*{c}{EPSR[2028]})" for s, c in zip(SC, "BCD")), None, None], "0.00",
     "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）"); r_ += 1
tset(r_, "comptp", "Comparable 目標價", [f"=B{T['pmed']}*B{T['feps']}", None, None], "#,##0", "同業中位數 × 上列 EPS"); r_ += 2
V.cell(r_, 1, "0.3 結論").font = BOLD; r_ += 1
hdr(V, r_, ["項目", "數值", "", "", "說明"]); T["h3"] = r_; r_ += 1
tset(r_, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], "#,##0", "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價（相對倍數調整後，見第 0.5 節）"); r_ += 1
tset(r_, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], "#,##0", "引用 0.2（未作相對倍數調整）"); r_ += 1
tset(r_, "cons", "市場共識目標價", ["=cons_tp", None, None], "#,##0", "stockanalysis 個股頁 Analyst price target 2,213（2026-10-22）"); r_ += 1
tset(r_, "upw", "上檔空間：加權目標價", [f"=B{T['tpw']}/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-30 收盤 1,265）"); r_ += 1
tset(r_, "upc", "上檔空間：Comparable 目標價", [f"=B{T['comp2']}/px-1", None, None], "0.0%"); r_ += 1
_u = f"B{T['upw']}"
V.cell(r_, 1, "上檔空間級距"); T["grade"] = r_
V[f"B{r_}"] = f'=IF({_u}>0.5,"大於 50%",IF({_u}>0.25,"25%–50%",IF({_u}>=0,"0%–25%",IF({_u}>=-0.25,"0%～−25%",IF({_u}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"E{r_}"] = "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）"; r_ += 1
V.cell(r_, 1, "投資建議"); T["rec"] = r_
V[f"B{r_}"] = f'=IF({_u}>0.5,"強力買進",IF({_u}>0.25,"買進",IF({_u}>=-0.25,"中立",IF({_u}>=-0.5,"放空","強力放空"))))'
V[f"E{r_}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空"; r_ += 1
tset(r_, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS 見 In 頁 cons_feps 與 Recon"); r_ += 1
tset(r_, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%"); r_ += 1
V[f"E{T['gmul']}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(r_, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）"); r_ += 1
PEROW = r_ + 3; PE10 = f"$B${PEROW + 1}:$K${PEROW + 1}"
tset(r_, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
     [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E 分布（下方列；2016–2017 年 EPS 未找到，共 8 年）的分位；低於最小值記 0%、高於最大值記 100%；只作檢查，不進目標價"); r_ += 1
V.cell(r_, 1, "目標價方法"); T["m"] = r_
V[f"B{r_}"] = ("DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年平均約 16%（含 2023 年谷底 7%，再投資率約 25%），FY2027–29 驅動因子、FY2030–32 收斂法＋永續；"
               "ROE 取三年平均使 DCF 偏低估（近 12 月 ROE 31%）。Comparable：台光電、聯茂、金像電前瞻 P/E 中位數（只有三家；金像電為下游 PCB，倍數較低），"
               "再乘相對倍數係數（第 0.5 節）。情境對應：A＝順利、B＝平平、C＝不如預期（Siliconomics 同字母顆數情境）。")
r_ += 1
CLOSE = [37.25, 83.80, 89.80, 148.00, 121.00, 102.00, 51.40, 126.00, 170.00, 494.00]
EPSY = [None, None, 7.46, 6.92, 6.67, 7.01, 4.69, 3.05, 9.56, 12.13]
V.cell(PEROW, 1, "十年年底 P/E（歷史分位檢查用；Yahoo Finance 12 月收盤 ÷ 當年 EPS〔HiStock〕；2016–2017 年 EPS 未找到）").font = BOLD
for j, yy in enumerate(range(2016, 2026)):
    V.cell(PEROW, 2 + j, str(yy))
    if EPSY[j] is not None:
        cc = V.cell(PEROW + 1, 2 + j, f"={CLOSE[j]}/{EPSY[j]}"); cc.number_format = "0.0"
V.cell(PEROW + 1, 1, "年底 P/E"); T["pe10h"] = PEROW; T["pe10"] = PEROW + 1
V.cell(PEROW + 2, 1, "核對：stockanalysis 年度 P/E 2021–2025＝14.58、10.97、41.46、17.79、40.96（與上列同量級，只作核對）")
for k_ in ("dcftp", "comptp", "tpw"): V[f"B{T[k_]}"].fill = KEY
V.cell(T["tpw"], 1).font = BOLD
defname("tp_final", f"Val!$B${T['tpw']}"); defname("tp_dcf", f"Val!$B${T['dcftp']}"); defname("tp_comp", f"Val!$B${T['comptp']}")
# 1. 價格與市場隱含倍數
v1 = PEROW + 5; T["s1"] = v1
V.cell(v1, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, v1 + 1, ["項目", "A", "B", "C", "說明"])
V.cell(v1 + 2, 1, "現價（2026-10-30 收盤）"); V[f"B{v1 + 2}"] = "=px"; V[f"B{v1 + 2}"].number_format = "#,##0"; T["px"] = v1 + 2
for k_, (lab, y) in enumerate((("現價 ÷ FY2026 EPS（錨點）", 2026), ("現價 ÷ FY2027 EPS（模型）", 2027), ("現價 ÷ FY2028 EPS（模型）", 2028), ("現價 ÷ FY2029 EPS（模型）", 2029))):
    rr = v1 + 3 + k_; V.cell(rr, 1, lab); T[f"pe{y}"] = rr
    for j, s in enumerate(SC):
        V[f"{'BCD'[j]}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{'BCD'[j]}{rr}"].number_format = "0.0"
rr = v1 + 7
V.cell(rr, 1, "近 12 月 EPS（4Q25＋1Q26～3Q26 揭露 EPS）"); V[f"B{rr}"] = "=3.76+4.36+8.02+12.88"; V[f"B{rr}"].number_format = "0.00"; T["ttm"] = rr
V.cell(rr + 1, 1, "  現價 ÷ 近 12 月 EPS"); V[f"B{rr + 1}"] = f"=px/B{rr}"; V[f"B{rr + 1}"].number_format = "0.0"; T["pettm"] = rr + 1
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens（一次改一個 In 頁低／高值、整本重算；數值由 tuc/sens_tuc.py 產生的 sens.json 寫入） ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次改一個輸入（In 頁敏感度低／高值），整本活頁簿以 LibreOffice 重算後的加權目標價（Val 第 0 節）與情境 A 的 FY2028 EPS；由 tuc/sens_tuc.py 產生 tuc/sens.json，建置時寫入（數值）"
SENS_F = os.path.join(HERE, "sens.json")
SJ = json.load(open(SENS_F, encoding="utf-8")) if (os.path.exists(SENS_F) and not OVR) else None
LBL = {x[0]: x[1] for x in INLIST}
SE["A2"] = ("相對基準＝案例目標價 ÷ 基準加權目標價 − 1；龍捲風依絕對值排序（HTML）。ai_25（2025 年高速材料占比）只影響 FY2025 欄的子列拆分，對目標價結構上恆為 0，仍列案例以示涵蓋。"
            "tp_g 案例同時移動 DCF 永續成長與 PL 收斂目標 fade_g（fade_g 以 tp_g 為公式）。")
hdr(SE, 3, ["驅動", "設定", "方向", "加權目標價", "相對基準", "鍵", "FY2028 EPS（情境 A）", "FY2028 EPS 相對基準", "說明"])
SE.cell(4, 1, "基準"); SE.cell(4, 2, "In 頁基準欄"); SE.cell(4, 3, "—"); SE.cell(4, 9, "本活頁簿重算值")
SE["D4"] = "=tp_final"; SE["G4"] = f"=PL!F{SUMR[('A', 'eps')]}"
SE["D4"].number_format = "#,##0"; SE["G4"].number_format = "0.00"; SE["E4"] = 0; SE["H4"] = 0
for c in "EH": SE[f"{c}4"].number_format = "0.0%"
SENS_KEYS = [x[0] for x in INLIST if x[3] is not None and x[5] is not None and x[3] != x[5]]
srow = 5
for k in SENS_KEYS:
    lo, base, hi = next((x[3], x[4], x[5]) for x in INLIST if x[0] == k)
    for side, val in (("低", lo), ("高", hi)):
        SE.cell(srow, 1, f"{LBL[k]}：{side}"); SE.cell(srow, 2, f"{base} → {val}"); SE.cell(srow, 3, side); SE.cell(srow, 6, k)
        if SJ and k in SJ:
            tp_, eps_ = SJ[k][side]
            SE.cell(srow, 4, tp_).number_format = "#,##0"; SE.cell(srow, 5, f"=D{srow}/D$4-1").number_format = "0.0%"
            SE.cell(srow, 7, eps_).number_format = "0.00"; SE.cell(srow, 8, f"=G{srow}/G$4-1").number_format = "0.0%"
            SE.cell(srow, 9, "重算值（sens_tuc.py）")
        else:
            SE.cell(srow, 9, "未產生（請先執行 sens_tuc.py）")
        srow += 1
SENS_LAST = srow - 1
SE.cell(srow + 1, 1, "讀法：產能受限的年度（情境 A 的 FY2027–28），需求面輸入（份額相對值、市場係數）往上調時營收被產能上限截斷，只有產能、組合與價格指數會改變營收；價格指數同時改變產能上限營收與高速材料毛利率（成本不變），所以對 EPS 的影響大於對營收。")
SENS_NOTE = srow + 1
SE.column_dimensions["A"].width = 60; SE.column_dimensions["B"].width = 18; SE.column_dimensions["I"].width = 30
for c in "CDEFGH": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識、券商估計與 Siliconomics 市場。只作事後對照，不為對上數字而調參數。FY2028 共識（最新）：未找到"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%" if not pt else "0.0000"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與券商（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs FactSet 共識（08-04）", "=c_rev26f", pl("rev", 2026), "FactSet（理財周刊 2026-09-10 轉引）", "3Q26 實際營收 179.5 億，共識在 3Q 前", fmt="#,##0.0")
rec("FY2026 EPS vs FactSet 共識（08-04）", "=c_eps26f", pl("eps", 2026), "同上", "")
rec("FY2026 EPS vs 美系券商（07-20）", "=b_eps26u", pl("eps", 2026), "經濟日報 2026-07-20", "")
rec("FY2027 營收 vs FactSet 共識（08-04）", "=c_rev27f", pl("rev", 2027), "FactSet", "", fmt="#,##0.0")
rec("FY2027 EPS vs FactSet 共識（08-04）", "=c_eps27f", pl("eps", 2027), "FactSet", "低 53.91、高 100.07")
rec("FY2027 EPS vs 美系券商（07-20）", "=b_eps27u", pl("eps", 2027), "經濟日報 2026-07-20", "")
rec("FY2028 EPS vs 美系券商（07-20）", "=b_eps28u", pl("eps", 2028), "經濟日報 2026-07-20", "單一券商、偏樂觀端")
rec("FY2028 EPS vs FactSet（04-20，單一預估）", "=c_eps28f", pl("eps", 2028), "鉅亨 2026-04-20", "期間早，只作量級參考")
RC.cell(rr_, 1, "B. 市場與份額（機制檢查）").font = BOLD; rr_ += 1
for y, j in ((2026, 0), (2027, 1), (2028, 2), (2029, 3)):
    rec(f"CY{y} AI 高速 CCL 市場 vs Siliconomics SIF_CCLMkt（US$B）", None, (lambda yy, jj: (lambda s: MKT[s][jj]))(y, j), SNAP, "下一欄為本模型（÷ mf）", fmt="#,##0.00", gap=False)
    for s_i, s in enumerate(SC):
        RC.cell(rr_ - 1, 4 + 2 * s_i, f"=PL!{FC[y]}{PB[s]['mku']}{'' if y == 2026 else '/mf'}-{chr(67 + 2 * s_i)}{rr_ - 1}").number_format = "0.0000"
rec("FY2026 台燿有效份額（需求面）", None, lambda s: f"=PL!D{PB[s]['sh0']}", "—", "台光電同期 45.9%（台光電成品 Units）", fmt="0.0%")
rec("FY2027 需求面 ÷ 產能上限", None, pl("util", 2027), "—", "大於 100%＝產能受限", fmt="0.0%")
rec("FY2028 需求面 ÷ 產能上限", None, pl("util", 2028), "—", "", fmt="0.0%")
rec("每張營收 FY2026（NT$；校準）", None, lambda s: f"=PL!D{PB[s]['rps0']}", "—", "FY2025：303.4 億 ÷（約 210 萬張 × 12）≈ NT$1,200；台光電 FY2026 約 NT$2,858", fmt="#,##0")
rec("其他等級毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "高速材料毛利率＝此值＋d_ai", fmt="0.0%")
rec("4Q26 EPS（模型）", None, lambda s: "=Q!E15", "—", "3Q26 實績 12.88", fmt="0.00")
RECON_LAST = rr_ - 1

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝台燿財報、月營收與法說，經 HiStock、nstock、富果轉述；券商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("財報：HiStock 損益表（季；公開資訊觀測站轉載；1Q25–3Q26 營收、毛利、營業利益、稅前、稅後）", HS_IS),
        ("財報：HiStock 每股盈餘（2018–2026 各季）", HS_EPS),
        ("財報：nstock（3Q26 EPS 12.88、毛利率 36.93%、營益率 31.02%；10-30 收盤）", NS),
        ("月營收：HiStock（2025-01～2026-09）", HS_REV),
        ("二手：優分析 2026-07-30（2Q26 營收 143.01 億、毛利率 29.72%、EPS 8.02；1H26 EPS 12.4；可轉債 100 億元）", UA),
        ("公司：2Q26 法說（富果 2026-07-30；損耗等級組合、泰國廠、2028–29 擴產 195 萬張）", FG0730),
        ("公司：4Q25 法說（富果 2026-03-12；ELL／SLL 36%、常熟與泰國各 +60 萬張）", FG0312),
        ("媒體：理財周刊 2026-09-10（集團月產能逾 220 萬張、市場估 2026／2027 年底 260／380 萬張；FactSet 08-04 共識）", MW),
        ("券商：經濟日報 2026-07-20（美系券商 2026–28 EPS 38.8／93.5／168.7）", UDN),
        ("券商：旺得富 2026-06-30（日系外資目標價 2,115；AI ASIC 專案 2026 年貢獻 5–6%）", WR),
        ("共識：鉅亨 2026-04-20（FactSet 2026–28 EPS、目標價 900）", CNY),
        ("資料商：stockanalysis（共識目標價 2,213、Forward PE 26.37、ROE、年度 P/E；同業 Forward PE）", SA),
        ("資料商：Yahoo Finance 月線收盤（2016–2025 年 12 月）", "https://finance.yahoo.com/quote/6274.TWO/history/"),
        ("市場層：Siliconomics 快照（部署 GW、每 GW CCL 金額、CCL 市場、匯率）；台光電成品 Units 頁（份額一致性檢查）", f"{SNAP}；{rel(EMC_F)}"),
        ("待補（人工）：10 月營收（約 11-10）；3Q26 法說（日期未找到）；財報附註的營運部門原文；2016–2017 年 EPS；FY2028 最新共識；Bloomberg 前瞻倍數歷史（已有 6274.TWO）", "未找到")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
OBS = [   # 重點觀察（數字取自重算後的 Val／PL／Sens；與 README.md 相同文字）
 "目標價 NT$1,828（DCF 1,343、Comparable 相對倍數調整後 2,312 各半），比現價 1,265（2026-10-30）高 44.5%，級距 25%–50%、投資建議買進；Comparable 單獨（市場主流同業法 2,486）為 +96.5%。"
 "上檔主因是 EPS：本模型 12 個月後前瞻 EPS 90.44 比共識前瞻 EPS 53.85（stockanalysis 2026-10-22，早於 3Q26 財報 EPS 12.88）高 67.9%；現價只隱含 13.4 倍情境 B 的 FY2028 EPS，同業前瞻 P/E 中位數 27.5 倍。DCF 低於 Comparable，因 ROE 取近三年平均 16%（含 2023 年谷底）。",
 "目標價最敏感的是 AI 高速 CCL 市場與價格：市場係數 ×0.8／×1.2 使加權目標價 −13.9%／+9.4%，2029 年售價指數 0.85 −10.5%、2028 年 0.90 −9.8%，2029 年份額相對值 0.90 −10.0%；DCF 權重 0.3／0.7 為 +10.6%／−10.6%（Sens 頁）。",
 "FY2027 三情境需求面都超過產能（117%），EPS 73.78 由月產能 310 萬張與組合、價格指數決定，比 FactSet 08-04 共識 70.47 高 4.7%；常熟、泰國各 60 萬張 2027Q3 開出是關鍵時點。台光電＋台燿需求面份額 2026–29 年合計 50–56%，低於 100%。下一個檢驗點：10 月營收（約 2026-11-10）與 3Q26 法說（日期未找到）。"]
lines = [(f"台燿（6274）個股模型 {VER}（AI 半導體 Project：載板與 PCB 次產業；高速 CCL）　{DATE}", True),
 (f"版本：{VER}（{DATE}，初版）。樣板 emc/（台光電）；行業變數讀 {SNAP}。", False),
 ("模型命題：台燿的高速材料（ELL／SLL，M7 以上）營收＝AI 高速 CCL 市場（Siliconomics：部署 GW × 每 GW CCL 金額）× 台燿有效份額 × 年均售價指數，但受月產能 × 每張營收限制（全產全銷）；其他等級（VLL／LL、非低損耗、特殊應用）另依成長率推。份額相對台光電一致（兩家合計 < 100%）。", False),
 ("核心問題：泰國、常熟、中山新產能（月產能約 235 萬 → 310 萬 → 430 萬 → 530 萬張）開出時，台燿在 AI 伺服器與交換器的 M7 以上份額與單價能撐多久；份額與價格回落時 EPS 還剩多少；現價隱含的路徑合不合理。", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)。預測期 FY2027–FY2029 由驅動因子推導，FY2030–FY2032 為收斂法。金額：十億新台幣；口徑：歸屬母公司淨利。", False),
 ("驅動因子：Siliconomics SIF_DeployGW_A／B／C × SIF_CCLperGW（＝SIF_CCLMkt_*）× 匯率 × 台燿有效份額（FY2026 以公司 ELL／SLL 占比校準）× 份額相對值（Assumed，待併入 Siliconomics）× 年均售價指數（Assumed，待併入 Siliconomics）；產能上限＝月產能 × 12 × 每張營收 × 組合指數 × 價格指數。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），FY2027–FY2032 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，推到 12 個月後）＋0.5 × Comparable 目標價（台光電、聯茂、金像電前瞻 P/E 中位數 × 12 個月後前瞻 EPS × 相對倍數係數）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("已知限制：(1) 高速材料（M7 以上）當作 AI 相關的代理，口徑未經公司確認；(2) 份額相對值、價格指數、組合指數皆 Assumed；(3) 官方營運部門以單一部門處理（財報附註原文未取得）；(4) 十年 P/E 只有 8 年（2016–2017 年 EPS 未找到）；(5) 10 月營收與 3Q26 法說未納入。", False),
 ("", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
RMX.apply(wb, "tuc")
apply_in_layout.apply(wb, VER, DATE)
# 份額一致性檢查（建置時；同一市場兩家需求面份額合計 < 100%）——以 Python 依 In 頁值重算 FY2026 校準份額（與 Excel 同式）
_B = {x[0]: OVR.get(x[0], x[4]) for x in INLIST}
_rev26 = 10.053966 + 14.301042 + 17.948807 * (2 + _B["q4g"])
_sh26 = _rev26 * _B["ai_26"] / (GW["A"][0] * PGW[0] / 1000 * _B["fx"])
for j, y in enumerate(BU):
    tot = EMC_SH[0][j] + _sh26 * (_B[f"rel_{y % 100}"] if y > 2026 else 1.0)
    assert tot < 1.0, ("台光電＋台燿份額超過 100%", y, tot)
# README：相對倍數與 In 頁版面工具的版本說明改寫為新模型用語；投資建議與重點觀察寫在工具之後（不被改寫）
for row_ in RD.iter_rows():
    for c_ in row_:
        if isinstance(c_.value, str) and c_.value.startswith("2026-10-29：加權目標價中的 Comparable"):
            c_.value = "相對倍數調整：加權目標價中的 Comparable 用相對倍數調整後的值（Val 第 0.5 節：同業前瞻倍數中位數 × 本公司 10 年相對同業的倍數比，Bloomberg 前瞻倍數；tools/relative_multiple.py）。「Comparable 目標價（市場主流同業法）」不調整。"
        elif isinstance(c_.value, str) and c_.value.startswith(f"{VER}（{DATE}）：依個股模型原則第 8 節"):
            c_.value = "In 頁版面（個股模型原則第 8 節）：D／F 欄＝敏感度低／高，只用於 Sens 頁、不進目標價；In 頁最後一節「情境驅動對照」把情境專屬輸入三情境並排（只引用上方各列，純展示）。"
rn = RD.max_row + 2; OBSROW = rn
RD.cell(rn, 1, f'="投資建議："&Val!B{T["rec"]}&"（依上檔空間固定規則；加權目標價上檔空間 "&TEXT(Val!B{T["upw"]},"+0.0%;−0.0%")&"）"'); rn += 1
for k_, o_ in enumerate(OBS, 1):
    RD.cell(rn, 1, f"重點觀察 {k_}：" + o_); rn += 1
wb.save(sys.argv[1])
if os.environ.get("TUC_ROWMAP"):
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "XT": XT, "XH": XH, "XR": XR, "T": T, "EPSR": EPSR, "CFR": CFR,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "UNOTE": UNOTE, "INROW": INROW, "OBSROW": OBSROW, "SENS_KEYS": SENS_KEYS, "nlines": len(lines)},
              open(os.environ["TUC_ROWMAP"], "w"), ensure_ascii=False, indent=1)
