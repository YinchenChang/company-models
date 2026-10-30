# Fabrinet（FN，NYSE）收支與估值模型 v0.1.2（AI 半導體 Project：網通次產業內複製——光模組與光元件代工，樣板 cohr/）
# 用法：python3 fabrinet/build_fabrinet.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度：FY(t)＝t−1 年 7 月～t 年 6 月（6 月最後一個星期五結束）。FY(t) 一半落在 CY(t−1) 下半年、一半落在 CY(t) 上半年。
#   市場模型是日曆年：FY2027 光模組＝CY2026 下半年（FY2027 Q1 指引推導＋Q2 假設）＋ hw × CY2027；FY2028＝（1 − hw）× CY2027 ＋ hw × CY2028；FY2029 同理。
# 金額：十億美元（US$B）；EPS：美元；口徑：公司 non-GAAP（排除股票薪酬、訴訟、資遣、重組、非上市股權重估、Pillar Two 補充稅）。
# FY2026（2026-06-26 結束）為完整實績，即錨點。營收類別依公司 FY2026 起的新分類：資料中心（光模組、DCI、HPC）、通訊基礎設施、車用工業與其他。
# 情境與行業一致（Andy 2026-10-25 13:13）：A／B／C 讀 tsmc 同字母顆數係數（＝markets/memory SCEN_UNITS）；公司特有變數（光模組份額、CPO 份額、毛利率）A 取高值、B 基準、C 低值。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
import cycle_position as CP
import valuation_rule as VR
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
import market as MEM

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx"); INL_F = only("innolight/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
NETR = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in ("  其中 CPO", "  其中可插拔", "AI IT 資本支出（加速器 × 每顆 IT capex）"):
        NETR[lab.strip()] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETR) == 3 and all(v is not None for k in NETR for v in NETR[k][0]), NETR.keys()
# 網通 CPO 三情境（markets/network v0.2.3；Andy 2026-10-25 13:13：技術替代屬行業情境，A＝替代快、C＝慢，公司讀同字母）
NETS = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    for s_ in ("A", "B", "C"):
        for k_, l_ in (("mpl", f"  其中可插拔（情境 {s_}）"), ("mcp", f"  其中 CPO（情境 {s_}）")):
            if r[0].value == l_:
                NETS[(k_, s_)] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETS) == 6 and all(v is not None for k in NETS for v in NETS[k][0]), NETS.keys()
NETR["其中可插拔"] = NETS[("mpl", "A")]; NETR["其中 CPO"] = NETS[("mcp", "A")]   # 情境 A 列（數值＝主表「其中可插拔」「其中 CPO」）
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
for s_ in SC:   # 情境與行業一致：tsmc 顆數係數必須等於 markets/memory 的 SCEN_UNITS（同一字母）
    assert abs(TS[f"um28_{s_}"][0] - MEM.SCEN_UNITS[s_][0]) < 1e-9 and abs(TS[f"um29_{s_}"][0] - MEM.SCEN_UNITS[s_][1]) < 1e-9, s_
INL = None
for r in openpyxl.load_workbook(INL_F, data_only=True)["Units"].iter_rows(min_row=4):
    if r[0].value == "中際旭創可插拔有效份額":
        INL = ([r[2 + j].value for j in range(4)], f"{rel(INL_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert INL and all(isinstance(v, float) for v in INL[0]), INL

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。情境 A 取公司特有變數的高值、B 取基準、C 取低值。SemiAnalysis 未使用。公司原文＝SEC 8-K 新聞稿與 10-K；資料中心內的光模組、DCI、HPC 拆分以法說會揭露金額推導"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
PR = "https://www.sec.gov/Archives/edgar/data/0001408710/000140871026000026/fn-2026811xex991q426.htm"
PR3 = "https://www.sec.gov/Archives/edgar/data/0001408710/000140871026000014/fn-2026504xex991q326.htm"
TENK = "https://www.sec.gov/Archives/edgar/data/0001408710/000140871026000028/fn-20260626.htm"
WB4 = "https://www.webull.com/news/15420495044125696"
NQ4 = "https://www.nasdaq.com/articles/fabrinet-q4-earnings-call-highlights"
P2 = "https://pulse2.com/fabrinet-data-center-revenue-jumps-68-to-669-million-and-becomes-51-of-quarterly-sales-as-dci-more-than-doubles/"
ZK3 = "https://www.tradingview.com/news/zacks:01f406102094b:0-fn-q3-earnings-beat-estimates-revenues-up-on-strong-telecom-hpc/"
AOL2 = "https://www.aol.com/articles/fabrinet-fn-q2-2026-earnings-232305108.html"
SA = "https://stockanalysis.com/stocks/fn/statistics/"
SAF = "https://stockanalysis.com/stocks/fn/forecast/"
SAFIN = "https://stockanalysis.com/stocks/fn/financials/"
SABS = "https://stockanalysis.com/stocks/fn/financials/balance-sheet/"
MB = "https://www.marketbeat.com/stocks/NYSE/FN/"
TF = "https://www.tickflow.io/stock/FN/forecast"
CMC = "https://companiesmarketcap.com/fabrinet/pe-ratio/"
MT = "https://www.macrotrends.net/stocks/charts/FN/fabrinet/pe-ratio"
MTM = "https://www.macrotrends.net/stocks/charts/FN/fabrinet/net-profit-margin"

sec("A. AI 光模組市場（讀 markets/network 成品；日曆年；情境 A 基準，B／C 依同字母顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mpl_{y}", f"AI 可插拔光模組市場 CY20{y}", "US$B", None, NETR["其中可插拔"][0][j], None, "Derived", NETR["其中可插拔"][1],
        "加速器顆數 × 每顆光學內容（scale-out 為主；2029 含 scale-up 光學）− CPO" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mcp_{y}", f"AI CPO 光學市場 CY20{y}", "US$B", None, NETR["其中 CPO"][0][j], None, "Derived", NETR["其中 CPO"][1],
        "共同封裝光學；情境 A 列（網通 v0.2.3 起 CPO 滲透分三情境，替代速度為行業變數，B、C 見 In 頁末，各 × 同字母顆數係數，不自訂）" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 CY20{y}", "US$B", None, NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][0][j], None, "Derived", NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][1], "Recon 用：Fabrinet 資料中心營收 ÷ AI IT 資本支出" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"inl_{y}", f"中際旭創可插拔有效份額 CY20{y}（份額口徑一致；加總檢查用）", "%", None, INL[0][j], None, "Derived", INL[1], "中際旭創模型為光模組份額權威；本模型 Fabrinet 份額與它相加不得超過 100%（Units 頁）" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝順利（基準）；B＝平平；C＝不如預期；與台積電、中際旭創、Coherent 模型相同（＝markets/memory SCEN_UNITS，建置時核對）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")

sec("B. Fabrinet 光模組份額與單價（光模組有效份額＝資料中心光模組與其他營收 ÷ 可插拔市場；代工營收含客供以外的料件；低＝情境 C、基準＝B、高＝A）")
inp("rel_27", "光模組有效份額 CY2027 ÷ CY2026（相對值）", "x", 0.80, 1.00, 1.15, "Assumed", WB4 + "：數個新光模組專案最快 FY2027 Q1 起量，一個商用（merchant）專案 12 月季開始、其他 CY2027 初（公司，Interested-party）",
    TENK + "：NVIDIA 占營收 FY2025 27.6% → FY2026 16.3%（客戶自建與多供應商分散；高值＝新客戶專案、低值＝NVIDIA 份額續降）")
inp("rel_28", "光模組有效份額 CY2028 ÷ CY2026", "x", 0.70, 1.00, 1.25, "Assumed", "同上；Building 10（約 200 萬平方呎）2027 年初完工，產能由約 55–58 億美元增至 85–93 億美元（公司）", "—")
inp("rel_29", "光模組有效份額 CY2029 ÷ CY2026", "x", 0.60, 0.95, 1.30, "Assumed", "同上", "—")
inp("cpo_27", *SIL.row9("cposh_fabrinet_27")[1:])
inp("cpo_28", *SIL.row9("cposh_fabrinet_28")[1:])
inp("cpo_29", *SIL.row9("cposh_fabrinet_29")[1:])
inp("pxr_27", "單價係數 CY2027（相對網通模型每顆內容；基準 1）", "x", 0.90, 1.0, 1.05, "Assumed", "網通模型每顆光學內容已含世代單價；本係數只在 Sens 動（降價快於模型）", "代工價格為成本加成；料件降價直接降營收（公司未揭露價格）")
inp("pxr_28", "單價係數 CY2028", "x", 0.80, 1.0, 1.05, "Assumed", "同上（與中際旭創、Coherent 模型相同）", "—")
inp("pxr_29", "單價係數 CY2029", "x", 0.72, 1.0, 1.05, "Assumed", "同上", "—")
inp("hw", "日曆年上半年占全年比重（光模組；會計年度換算用）", "%", 0.42, 0.46, 0.50, "Assumed", "CY2026 推導：上半年（FY2026 Q3＋Q4）÷ 全年約 46%（Q 頁）", "與 Coherent 模型相同；Sens 動 42%／50%")
inp("q2_tg", "FY2027 Q2（2026-12 季）光模組季增", "%", 0.03, 0.08, 0.15, "Assumed", WB4 + "：商用光模組專案 12 月季開始（公司）", NQ4 + "：資料中心 FY2027 Q1 預期光模組、DCI、HPC 全面成長")

sec("C. FY2026 實績與新營收類別（公司 FY2026 起改報資料中心、通訊基礎設施、車用工業與其他三類）")
inp("q3_rev", "FY2026 Q3（2026-03 季）營收", "US$B", None, 1.2143, None, "Verified", PR3 + "：1,214.3 百萬美元", ZK3 + "：12.14 億美元（年增 39.3%）")
inp("q4_rev", "FY2026 Q4（2026-06 季）營收", "US$B", None, 1.315788, None, "Verified", PR + "：1,315,788 千美元（年增 45%）", NQ4 + "：13.16 億美元")
inp("q3_dc", "FY2026 Q3 資料中心營收（新分類）", "US$B", None, 0.592, None, "Interested-party", P2 + "：5.920 億美元（Q4 季增 13% 的基期；公司簡報經媒體轉述）", NQ4 + "：Q4 季增 13%（669 ÷ 1.13 ≈ 592）")
inp("q4_dc", "FY2026 Q4 資料中心營收", "US$B", None, 0.6693, None, "Interested-party", P2 + "：6.693 億美元（年增 68%、占 51%）", WB4 + "：6.69 億美元")
inp("q4_ci", "FY2026 Q4 通訊基礎設施營收", "US$B", None, 0.4129, None, "Interested-party", P2 + "：4.129 億美元（年增約 40%）", NQ4 + "：4.13 億美元，季增 1%")
inp("q4_aio", "FY2026 Q4 車用、工業與其他營收", "US$B", None, 0.2336, None, "Interested-party", P2 + "：2.336 億美元（年增約 8%）", NQ4 + "：2.34 億美元，季增 9%（EV 充電、LiDAR）")
inp("q4_cig", "FY2026 Q4 通訊基礎設施季增（Q3 回推用）", "%", None, 0.01, None, "Interested-party", NQ4 + "：季增 1%（公司）", WB4 + "：同值")
for k, lab, v, n1, n2 in (("dci_q1", "FY2026 Q1", 0.1381, "1.381 億美元", "—"), ("dci_q2", "FY2026 Q2", 0.1422, "1.422 億美元", AOL2 + "：1.42 億美元，年增 42%"),
                          ("dci_q3", "FY2026 Q3", 0.1969, "1.969 億美元", ZK3 + "：1.969 億美元，年增 90.4%"), ("dci_q4", "FY2026 Q4", 0.2559, "2.559 億美元", WB4 + "：DCI 年化約 10 億美元")):
    inp(k, f"{lab} DCI（資料中心互連）營收", "US$B", None, v, None, "Interested-party", P2 + "：" + n1 + "（法說會）", n2)
for k, lab, v, n1, n2 in (("hpc_q1", "FY2026 Q1", 0.0154, "0.154 億美元", AOL2 + "：Q1 約 0.15 億美元"), ("hpc_q2", "FY2026 Q2", 0.086, "0.86 億美元", "第二來源：未找到"),
                          ("hpc_q3", "FY2026 Q3", 0.1067, "1.067 億美元", "第二來源：未找到"), ("hpc_q4", "FY2026 Q4", 0.118, "1.18 億美元", "第二來源：未找到（法說會只說優於預期）")):
    inp(k, f"{lab} HPC（高效能運算）營收", "US$B", None, v, None, "Interested-party", ("Pulse 2.0" if k in ("hpc_q1", "hpc_q4") else ("Motley Fool 逐字稿（AOL）" if k == "hpc_q2" else "Zacks（TradingView）")) + "（網址見 Sources）：" + n1, n2)
inp("dci_y2", "DCI 年增 FY2026 Q2（FY2025 回推用）", "%", None, 0.42, None, "Interested-party", AOL2 + "：年增 42%", "—")
inp("dci_y3", "DCI 年增 FY2026 Q3（FY2025 回推用）", "%", None, 0.904, None, "Interested-party", ZK3 + "：年增 90.4%", "—")
inp("dci_q4_25", "FY2025 Q4 DCI 營收", "US$B", None, 0.107, None, "Interested-party", P2 + "：1.07 億美元", "—")
inp("dci_q1_25", "FY2025 Q1 DCI 營收", "US$B", 0.08, 0.10, 0.11, "Assumed", "未找到；假設與 FY2025 Q2 回推值（約 1.0 億）相近", "只影響 FY2025 拆分")
inp("nv26", "NVIDIA 占營收 FY2026", "%", None, 0.163, None, "Verified", TENK + "：16.3%（FY2025 27.6%）", WB4 + "：16%（Cisco 20%、Nokia 11%、Amazon 11%）")

sec("D. FY2027 Q1 指引與 Q2 假設（＝CY2026 下半年）")
inp("q1_rev", "FY2027 Q1（2026-09 季）營收指引中點", "US$B", 1.375, 1.40, 1.425, "Interested-party", PR + "：13.75–14.25 億美元（公司指引）", WB4 + "：中點年增 43%")
inp("q1_eps", "FY2027 Q1 non-GAAP EPS 指引中點", "US$", 4.10, 4.175, 4.25, "Interested-party", PR + "：4.10–4.25 美元（公司指引）", "GAAP 3.39–3.54；差額含 Pillar Two 0.40 美元")
inp("q1_gm", "FY2027 Q1 non-GAAP 毛利率", "%", 0.120, 0.122, 0.124, "Assumed", PR + "：FY2026 Q4 12.2%、全年 12.2%（公司不指引毛利率）", WB4 + "：Q1 季節性費用增加，短暫壓利潤率（公司）")
inp("q1_opx", "FY2027 Q1 non-GAAP 營業費用", "US$B", 0.018, 0.0195, 0.021, "Assumed", "FY2026 Q4 非 GAAP 營業費用約 0.175 億美元（毛利 − 營業利益）", "Q1 調薪季節性上升（公司）")
inp("q_dcig", "DCI 季增（FY2027 Q1、Q2）", "%", 0.05, 0.10, 0.15, "Assumed", "FY2026 Q4 DCI 季增 30%（推導）", NQ4 + "：DCI 是 Q4 最大成長來源")
inp("q_hpg", "HPC 季增（FY2027 Q1、Q2）", "%", 0.0, 0.10, 0.20, "Assumed", "FY2026 Q4 季增約 11%（推導）", AOL2 + "：Fabrinet 在 AWS 專案是第二供應商，非獨家（公司）")
inp("q_cig", "通訊基礎設施季增（FY2027 Q1、Q2）", "%", 0.0, 0.03, 0.05, "Assumed", NQ4 + "：預期健康成長（公司）", "—")
inp("q_aig", "車用、工業與其他季增（FY2027 Q1、Q2）", "%", 0.0, 0.02, 0.05, "Assumed", NQ4 + "：預期成長，未給數字（公司）", "—")

sec("E. DCI、HPC、通訊基礎設施、車用工業（FY2027＝上半年 ×（2＋下半年 ÷ 上半年 − 1）；FY2028 起年增；DCI、HPC 另乘同字母 AI 情境係數）")
inp("h2_dci", "DCI 下半年 ÷ 上半年 − 1（FY2027）", "%", 0.0, 0.12, 0.25, "Assumed", WB4 + "：DCI 年化約 10 億美元（Q4）", "—")
inp("g_dci_28", "DCI 年增 FY2028（基準路徑）", "%", 0.05, 0.20, 0.30, "Assumed", "scale-across 資料中心互連（800G ZR／ZR+）", "—")
inp("g_dci_29", "DCI 年增 FY2029（基準路徑）", "%", 0.0, 0.10, 0.20, "Assumed", "同上", "—")
inp("h2_hpc", "HPC 下半年 ÷ 上半年 − 1（FY2027）", "%", -0.10, 0.08, 0.20, "Assumed", "AWS 專案（公司未具名）第二供應商；量能波動大", "—")
inp("g_hpc_28", "HPC 年增 FY2028（基準路徑）", "%", -0.10, 0.15, 0.30, "Assumed", "同上", "—")
inp("g_hpc_29", "HPC 年增 FY2029（基準路徑）", "%", -0.10, 0.10, 0.20, "Assumed", "同上", "—")
inp("h2_ci", "通訊基礎設施下半年 ÷ 上半年 − 1（FY2027）", "%", 0.0, 0.04, 0.06, "Assumed", NQ4 + "：電信系統、衛星通訊、電信元件全面成長（公司）", "—")
inp("g_ci_28", "通訊基礎設施年增 FY2028", "%", 0.0, 0.06, 0.10, "Assumed", "FY2026 年增約 47%（10-K 比重推導）；高基期後放緩", "—")
inp("g_ci_29", "通訊基礎設施年增 FY2029", "%", -0.03, 0.05, 0.08, "Assumed", "同上", "—")
inp("h2_aio", "車用、工業與其他下半年 ÷ 上半年 − 1（FY2027）", "%", -0.02, 0.02, 0.05, "Assumed", NQ4 + "：EV 充電與 LiDAR（公司）", "—")
inp("g_aio_28", "車用、工業與其他年增 FY2028", "%", 0.0, 0.05, 0.08, "Assumed", "FY2026 年增約 10%（10-K 比重推導）", "—")
inp("g_aio_29", "車用、工業與其他年增 FY2029", "%", 0.0, 0.05, 0.08, "Assumed", "同上", "—")

sec("F. 毛利率、費用、稅率、股數（non-GAAP；毛利率低＝情境 C、基準＝B、高＝A）")
inp("gm_27", "non-GAAP 毛利率 FY2027", "%", 0.118, 0.123, 0.127, "Assumed", PR + "：FY2026 12.2%、FY2025 12.4%", WB4 + "：Q4 12.2%（季增 10bp、年減 30bp）；規模帶來營業槓桿（公司）")
inp("gm_28", "non-GAAP 毛利率 FY2028", "%", 0.115, 0.123, 0.128, "Assumed", "同上；新廠爬坡期壓毛利（低值）", "—")
inp("gm_29", "non-GAAP 毛利率 FY2029", "%", 0.113, 0.123, 0.128, "Assumed", "同上", "—")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.10, "=Q!B24", 0.60, "Derived", "FY2026 non-GAAP 營業費用年增 ÷ 營收年增（Q 頁）", WB4 + "：營業費用約占營收 1.3%（Q4）")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.002, 0.0, 0.005, "Assumed", "無", "Sens 用 +0.5 個百分點")
inp("int_f", "non-GAAP 業外淨收益（利息、匯兌、其他；FY2027 起，每年）", "US$B", 0.020, 0.030, 0.040, "Assumed", PR + "：FY2026 利息收入 3,240 萬美元；現金與短期投資 8.75 億美元（2026-06）", "資本支出（FY2026 2.53 億美元）使現金下降")
inp("tx", "non-GAAP 稅率（不含 Pillar Two 補充稅）", "%", 0.04, 0.05, 0.06, "Assumed", "FY2026 non-GAAP 約 4.6%（Q 頁推導：GAAP 稅 8,209 萬 − Pillar Two 5,745 萬）", "泰國 BOI 租稅優惠；明細：未找到")
inp("p2", "Pillar Two 補充稅占稅前（Sens 用；non-GAAP 排除）", "pt", None, 0.09, None, "Derived", PR + "：FY2027 Q1 指引排除 Pillar Two 每股 0.40 美元（約 1,450 萬美元／季）", TENK + "：泰國 UTPR 已立法；FY2026 認列 5,745 萬美元（現金尚未支付）")
inp("sh_q1", "FY2027 Q1 稀釋股數（指引）", "十億股", None, 0.0363, None, "Interested-party", PR + "：約 3,630 萬股（公司指引）", TENK + "：普通股 35,834,864 股（2026-08-07）")
inp("sh_g", "股數年變動", "%", -0.01, 0.0, 0.01, "Assumed", PR + "：FY2026 回購 520 萬美元（FY2025 1.26 億）；剩餘授權約 1.69 億", "股票薪酬稀釋約抵回購")
inp("mf", "AI 市場係數（Sens 用；CY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 CY2027 可插拔光模組 +84%", "±20%")

sec("G. FY2025、FY2026 實績（non-GAAP；US$B）")
for k, lab, v25, v26, n in (
        ("rev", "營收", 3.419327, 4.641097, "3,419,327／4,641,097 千美元"),
        ("dcp", "資料中心占營收（新分類）", 0.462, 0.479, "46.2%／47.9%（10-K）"),
        ("cip", "通訊基礎設施占營收", 0.307, 0.333, "30.7%／33.3%（10-K）"),
        ("aiop", "車用、工業與其他占營收", 0.231, 0.188, "23.1%／18.8%（10-K）"),
        ("gp", "non-GAAP 毛利", 0.423805, 0.56797, "423,805／567,970"),
        ("op", "non-GAAP 營業利益", 0.360712, 0.500063, "360,712／500,063"),
        ("tax", "non-GAAP 所得稅（GAAP 稅 − Pillar Two）", 0.022653, "=0.082086-0.057447", "22,653／82,086 − 57,447"),
        ("ni", "non-GAAP 淨利", 0.368792, 0.510907, "368,792／510,907"),
        ("eps", "non-GAAP 稀釋 EPS", 10.17, 14.09, "10.17／14.09 美元"),
        ("sh", "稀釋股數", 0.036267, 0.036252, "36,267／36,252 千股")):
    unit = "US$" if k == "eps" else ("十億股" if k == "sh" else ("%" if k.endswith("p") else "US$B"))
    src1 = (TENK if k.endswith("p") else PR) + "：" + n.split("／")[0]
    src2 = (TENK if k.endswith("p") else PR) + "：" + n.split("／")[-1]
    inp(f"a25_{k}_in", f"FY2025 {lab}", unit, None, v25, None, "Verified", src1, SAFIN + "：營收 3,419 百萬美元（stockanalysis）" if k == "rev" else "—")
    inp(f"a26_{k}_in", f"FY2026 {lab}", unit, None, v26, None, "Verified", src2, SAFIN + "：營收 4,641 百萬美元" if k == "rev" else (WB4 + "：14.09 美元（法說會）" if k == "eps" else (NQ4 + "：Q4 資料中心 51%" if k == "dcp" else "—")))

sec("H. 共識、市場資料（只作 Recon 對帳，不作校準）")
inp("cons_rev27", "共識 FY2027 營收", "US$B", None, 6.11, None, "Interested-party", SAF + "：61.1 億美元（7 位分析師，2026-10-21 更新）", TF + "：54.6 億美元（2026-10-24；疑似未更新）")
inp("cons_rev28", "共識 FY2028 營收", "US$B", None, 6.58, None, "Interested-party", "tickflow（網址見 Sources）：65.8 億美元（3 位，2026-10-24；FY2027 值偏低，疑似未更新）", "第二來源：未找到")
inp("cons_eps27", "共識 FY2027 EPS", "US$", None, 18.23, None, "Interested-party", SAF + "：18.23（7 位分析師，區間 17.48–19.20）", MB + "：17.50（本年度，2026-10-23）")
inp("cons_eps27b", "共識 FY2027 EPS（第二來源）", "US$", None, 17.50, None, "Interested-party", MB + "：17.50（MarketBeat 未標年度，以本年度判讀）", "—")
inp("cons_eps28", "共識 FY2028 EPS", "US$", None, 21.27, None, "Interested-party", MB + "：21.27（下一年度，2026-10-23）", TF + "：20.14（2 位）")
inp("cons_eps28b", "共識 FY2028 EPS（第二來源）", "US$", None, 20.14, None, "Interested-party", TF + "：20.14（2 位分析師，2026-10-24）", "—")
inp("cap_b10", "產能（現有＋Pinehurst 轉換＋Building 10；營收額）", "US$B", 8.5, 8.9, 9.3, "Interested-party", WB4 + "：現有約 55–58 億，Building 10 增 30–35 億，合計 85–93 億美元（公司）", "Building 10 2027 年初完工")
inp("cap_tot", "長期產能目標（再加兩座 Chonburi 廠等）", "US$B", 12.5, 13.25, 14.0, "Interested-party", WB4 + "：未來數年 125–140 億美元（公司）", "—")

sec("I. 價格與估值")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 314.94, None, "Interested-party", SA + "：314.94（−0.86%）", MB + "：315.09（10-23）")
inp("px_hi", "52 週高點", "US$", None, 748.89, None, "Interested-party", "MarketBeat（網址見 Sources）：52 週區間 305.63–748.89", "第二來源：未找到（stockanalysis 只列 52 週變動 −21.3%）")
for s, w in zip(SC, (0.3, 0.5, 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")

sec("K. 目標價（DCF＋Comparable）與收斂法參數（黃底，Assumed，Andy 可改）")
for y, ni_, e0, e1 in ((24, 0.29618, 1.469, 1.746), (25, 0.332527, 1.746, 1.981812), (26, 0.473027, 1.981812, 2.454256)):
    inp(f"roe_{y}", f"GAAP ROE FY20{y}＝淨利 ÷ 平均股東權益", "%", None, f"={ni_}/(({e0}+{e1})/2)", None, "Derived",
        (PR if y >= 25 else SAFIN) + f"：淨利 {ni_ * 1000:,.1f} 百萬美元", SABS + f"：股東權益 {e0 * 1000:,.0f} → {e1 * 1000:,.0f} 百萬美元")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；各市場同一值）", "DCF 與 Sens")
inp("tp_g", "長期成長率（DCF 永續；收斂法情境 B 的 FY2032 營收年增）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "—")
inp("tp_roe", "ROE（近三年平均；DCF 再投資率＝tp_g ÷ tp_roe）", "%", None, "=AVERAGE(roe_24,roe_25,roe_26)", None, "Derived", "In 頁 roe_24–roe_26（GAAP；non-GAAP 調整項小，FY2026 重估利益與 Pillar Two 約相抵）", SA + "：ROE 21.33%（TTM）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（工作單 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("nim10", "十年平均淨利率（2016–2025；GAAP）", "%", None, "=AVERAGE(0.0736,0.0622,0.0714,0.0731,0.0726,0.0824,0.0949,0.0934,0.1027,0.0972)", None, "Derived",
    MTM + "：2016–2023 年底 TTM 淨利率 7.36%、6.22%、7.14%、7.31%、7.26%、8.24%、9.49%、9.34%", SAFIN + "：FY2024 10.27%、FY2025 9.72%（GAAP；比 non-GAAP 低約 1pt）")
inp("fade_gA", "收斂法：情境 A 於 FY2032 的營收年增（＝tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法；FY2030 起由 FY2029 年增線性收斂", "FY2030 起為收斂法，非驅動因子推導")
inp("fade_gB", "收斂法：情境 B 於 FY2032 的營收年增（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "收斂法：情境 C 於 FY2032 的營收年增（＝tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "收斂法：情境 A 的中段淨利率（＝情境 A FY2029 淨利率）", "%", None, None, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "收斂法：情境 B 的中段淨利率（＝情境 B FY2029 與十年平均的中點）", "%", None, None, None, "Derived", "（PL 頁情境 B FY2029＋nim10）÷ 2", "十年平均為 GAAP 口徑，略低於 non-GAAP（偏保守）")
inp("fade_mC", "收斂法：情境 C 的中段淨利率（＝十年平均）", "%", None, None, None, "Derived", "＝nim10", "Andy 可改")
inp("cons_tp", "市場共識目標價（賣方平均）", 'US$', None, 687.78, None, "Interested-party", SAF + "：9 位分析師平均 687.78（中位數 739、區間 431–800；2026-10-21 更新）", MB + "：10 位平均 650.63（區間 431–800）")
inp("cons_feps", '共識 12 個月後前瞻 EPS（≈ FY2028；FY2029 共識未找到）', 'US$', None, '=cons_eps28', None, "Derived", 'MarketBeat FY2028 共識 21.27（In 頁 cons_eps28）；本模型前瞻 EPS 含 4／12 的 FY2029，口徑略前', "共識只放 Recon 與本列，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
INROW = {}
rr = 4
for srow in rows:
    if srow[0] != "sec": INROW[srow[0]] = rr
    rr += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績、FY2026 季度與 FY2027 上半年 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與季度拆分（十億美元；EPS 為美元；股數為十億股；non-GAAP）。藍字＝輸入（In 頁）；黑字＝公式。FY(t)＝t−1 年 7 月～t 年 6 月"
Q["A2"] = "資料中心內的光模組、DCI、HPC 金額只在法說會揭露：光模組與其他＝資料中心 − DCI − HPC。FY2026 Q3 通訊基礎設施＝Q4 ÷（1＋季增 1%），車用工業＝營收 − 其他兩類。FY2027 Q1：各類別依季增假設，資料中心＝營收指引中點 − 通訊 − 車用工業；Q2 依季增假設。FY2027 Q1 季報（預估 2026-11-02）公布後以實績取代"
hdr(Q, 3, ["項目", "FY2025A", "FY2026A", "FY2026 Q3", "FY2026 Q4", "FY2027 Q1（指引）", "FY2027 Q2（假設）", "標記", "說明"])
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", note=""):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, note)
qrow(4, "營收", {"B": "=a25_rev_in", "C": "=a26_rev_in", "D": "=q3_rev", "E": "=q4_rev", "F": "=q1_rev", "G": "=G5+G9+G10"}, tag="Verified／Interested-party（指引）")
qrow(5, "資料中心", {"B": "=a25_rev_in*a25_dcp_in", "C": "=a26_rev_in*a26_dcp_in", "D": "=q3_dc", "E": "=q4_dc", "F": "=F4-F9-F10", "G": "=G6+G7+G8"}, tag="Verified（比重）／Interested-party（季）",
     note="FY2025、FY2026＝營收 × 10-K 比重；FY2027 Q1＝指引中點 − 通訊 − 車用工業")
qrow(6, "  光模組與其他（資料中心 − DCI − HPC）", {c: f"={c}5-{c}7-{c}8" for c in "BCDEF"} | {"G": "=F6*(1+q2_tg)"}, tag="Derived／Assumed",
     note="NVIDIA 光模組代工與其他客戶光模組；Q2 依季增假設")
qrow(7, "  DCI（資料中心互連）", {"B": "=dci_q1_25+dci_q2/(1+dci_y2)+dci_q3/(1+dci_y3)+dci_q4_25", "C": "=dci_q1+dci_q2+dci_q3+dci_q4", "D": "=dci_q3", "E": "=dci_q4", "F": "=E7*(1+q_dcig)", "G": "=F7*(1+q_dcig)"},
     tag="Interested-party／Assumed", note="FY2025＝Q1 假設＋Q2、Q3 由年增率回推＋Q4 實績")
qrow(8, "  HPC（高效能運算）", {"B": "=0", "C": "=hpc_q1+hpc_q2+hpc_q3+hpc_q4", "D": "=hpc_q3", "E": "=hpc_q4", "F": "=E8*(1+q_hpg)", "G": "=F8*(1+q_hpg)"},
     tag="Interested-party／Assumed", note="FY2025 視為 0（HPC 自 FY2026 Q1 起量，約 0.15 億美元）")
qrow(9, "通訊基礎設施（電信系統、衛星通訊、電信元件）", {"B": "=a25_rev_in*a25_cip_in", "C": "=a26_rev_in*a26_cip_in", "D": "=q4_ci/(1+q4_cig)", "E": "=q4_ci", "F": "=E9*(1+q_cig)", "G": "=F9*(1+q_cig)"}, tag="Verified／Derived")
qrow(10, "車用、工業與其他（EV 充電、LiDAR、工業雷射、醫療）", {"B": "=a25_rev_in*a25_aiop_in", "C": "=a26_rev_in*a26_aiop_in", "D": "=D4-D5-D9", "E": "=q4_aio", "F": "=E10*(1+q_aig)", "G": "=F10*(1+q_aig)"}, tag="Verified／Derived")
qrow(11, "non-GAAP 毛利", {"B": "=a25_gp_in", "C": "=a26_gp_in", "F": "=F4*q1_gm"}, tag="Verified")
qrow(12, "non-GAAP 毛利率", {c: f"={c}11/{c}4" for c in "BCF"}, fmt="0.0%", tag="Derived")
qrow(13, "non-GAAP 營業費用（毛利 − 營業利益）", {"B": "=a25_gp_in-a25_op_in", "C": "=a26_gp_in-a26_op_in", "F": "=q1_opx"}, tag="Derived")
qrow(14, "non-GAAP 營業利益", {c: f"={c}11-{c}13" for c in "BCF"}, tag="Derived")
qrow(15, "non-GAAP 業外淨收益（淨利＋稅 − 營業利益）", {"B": "=a25_ni_in+a25_tax_in-a25_op_in", "C": "=a26_ni_in+a26_tax_in-a26_op_in", "F": "=int_f/4"}, tag="Derived", note="利息收入、匯兌、其他；FY2026 排除非上市股權重估利益 5,674 萬美元")
qrow(16, "non-GAAP 稅率", {"B": "=a25_tax_in/(B14+B15)", "C": "=a26_tax_in/(C14+C15)", "F": "=tx"}, fmt="0.0%", tag="Derived")
qrow(17, "non-GAAP 淨利", {"B": "=a25_ni_in", "C": "=a26_ni_in", "F": "=(F14+F15)*(1-F16)"}, tag="Verified／Derived")
qrow(18, "稀釋股數", {"B": "=a25_sh_in", "C": "=a26_sh_in", "F": "=sh_q1"}, fmt="0.0000", tag="Verified")
qrow(19, "non-GAAP 稀釋 EPS", {c: f"={c}17/{c}18" for c in "BCF"}, fmt="0.00", tag="Derived", note="FY2027 Q1 指引中點 4.175")
Q.cell(21, 1, "推導量").font = BOLD
Q.cell(22, 1, "FY2026 營收年增"); Q["B22"] = "=C4/B4-1"
Q.cell(23, 1, "FY2026 營業費用年增"); Q["B23"] = "=C13/B13-1"
Q.cell(24, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B24"] = "=B23/B22"
Q.cell(25, 1, "CY2026 光模組與其他（FY2026 Q3＋Q4＋FY2027 Q1＋Q2；份額校準基礎）"); Q["B25"] = "=D6+E6+F6+G6"
Q.cell(26, 1, "CY2026 下半年光模組與其他（＝FY2027 上半年）"); Q["B26"] = "=F6+G6"
Q.cell(27, 1, "CY2026 上半年光模組 ÷ 全年（hw 的依據）"); Q["B27"] = "=(D6+E6)/B25"
Q.cell(28, 1, "FY2027 Q1 EPS（模型，指引中點）÷ 公司 EPS 指引中點 − 1"); Q["B28"] = "=F19/q1_eps-1"
Q.cell(29, 1, "FY2026 NVIDIA 營收（16.3% × 營收）÷ 光模組與其他"); Q["B29"] = "=nv26*C4/C6"
Q.cell(30, 1, "FY2026 資料中心 ÷ 營收"); Q["B30"] = "=C5/C4"
Q.cell(31, 1, "FY2027 上半年 DCI（Q1＋Q2）"); Q["B31"] = "=F7+G7"
Q.cell(32, 1, "FY2027 上半年 HPC"); Q["B32"] = "=F8+G8"
Q.cell(33, 1, "FY2027 上半年通訊基礎設施"); Q["B33"] = "=F9+G9"
Q.cell(34, 1, "FY2027 上半年車用、工業與其他"); Q["B34"] = "=F10+G10"
for rr_, f in ((22, "0.0%"), (23, "0.0%"), (24, "0.00"), (25, "#,##0.000"), (26, "#,##0.000"), (27, "0.0%"), (28, "0.0%"), (29, "0%"), (30, "0.0%"), (31, "#,##0.000"), (32, "#,##0.000"), (33, "#,##0.000"), (34, "#,##0.000")): Q[f"B{rr_}"].number_format = f
for nm, rr_ in (("rev", 4), ("dc", 5), ("tr", 6), ("dci", 7), ("hpc", 8), ("ci", 9), ("aio", 10), ("gp", 11), ("opx", 13), ("op", 14), ("oth", 15), ("ni", 17), ("sh", 18), ("eps", 19)):
    defname(f"a25_{nm}", f"Q!$B${rr_}"); defname(f"a26_{nm}", f"Q!$C${rr_}")
defname("tr_cy26", "Q!$B$25"); defname("h1_27", "Q!$B$26")
for nm, rr_ in (("dci", 31), ("hpc", 32), ("ci", 33), ("aio", 34)): defname(f"h1{nm}", f"Q!$B${rr_}")
Q.column_dimensions["A"].width = 60
for c in "BCDEFG": Q.column_dimensions[c].width = 15
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 90
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
mk = lambda b, y, s: f"{b}_{y % 100}" if s == "A" else f"{b}{s}_{y % 100}"   # 網通 CPO 三情境：情境 A 讀 A 節，B、C 讀同字母列
DEFP = dict(pxr="pxr", mf="1", hw="hw", gd="g_dci", gh="g_hpc", h2d="h2_dci", h2h="h2_hpc", opxa="0", tx="tx")
SCDEF = {"A": dict(rel="RHI", cpo="CPHI", gm="GMHI"), "B": dict(rel="rel", cpo="cpo", gm="gm"), "C": dict(rel="RLO", cpo="CPLO", gm="GMLO")}
def engine(ws, r0, title, scn="B", **over):
    P = dict(DEFP); P.update(SCDEF[scn]); P.update(over); s = scn; hw = P["hw"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026A（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    F = BU[1:]
    sec_("A. 日曆年 AI 光學（本節 D–G 欄＝CY2026–CY2029；D 欄＝CY2026 校準年）")
    row("mpl", "AI 可插拔光模組市場 × 情境顆數係數（日曆年）", "US$B", {c_(y): f"={mk('mpl', y, s)}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU}, note="本節欄位為日曆年")
    row("mcp", "AI CPO 光學市場 × 情境顆數係數（日曆年）", "US$B", {c_(y): f"={mk('mcp', y, s)}*{umx(y, s)}" for y in BU})
    row("sh0", "Fabrinet 光模組有效份額（CY2026 校準）＝ CY2026 光模組與其他 ÷ 可插拔市場", "%", {"D": f"=tr_cy26/D{R['mpl']}"}, fmt="0.0%",
        note="Derived：CY2026 CPO 營收視為 0；代工營收含料件，屬「有效」份額")
    row("rel", "份額相對值（÷ CY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in F}}, fmt="0.00")
    row("sh", "光模組有效份額（日曆年）", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("pxr", "單價係數（相對網通模型）", "x", {"D": "=1", **{c_(y): f"={P['pxr']}_{y % 100}" for y in F}}, fmt="0.00")
    row("cps", "Fabrinet 占 CPO 光學金額", "%", {"D": "=0", **{c_(y): f"={P['cpo']}_{y % 100}" for y in F}}, fmt="0.0%")
    row("tcy", "光模組與其他營收（日曆年）", "US$B", {"D": "=tr_cy26", **{c_(y): f"={c_(y)}{R['mpl']}*{c_(y)}{R['sh']}*{c_(y)}{R['pxr']}+{c_(y)}{R['mcp']}*{c_(y)}{R['cps']}" for y in F}})
    sec_("B. 會計年度營收（本節起 C–G 欄＝FY2025–FY2029）")
    row("uf", "AI 情境係數（會計年度；DCI、HPC 用）", "x", {"C": "=1", "D": "=1", "E": "=1", "F": f"=(1-{hw})+{hw}*{umx(2028, s)}", "G": f"=(1-{hw})*{umx(2028, s)}+{hw}*{umx(2029, s)}"}, fmt="0.000",
        note="同字母顆數係數依日曆年上半年比重換算；FY2027 為 1")
    row("tr", "  光模組與其他", "US$B", {"C": "=a25_tr", "D": "=a26_tr",
        "E": f"=h1_27+{hw}*E{R['tcy']}", "F": f"=(1-{hw})*E{R['tcy']}+{hw}*F{R['tcy']}", "G": f"=(1-{hw})*F{R['tcy']}+{hw}*G{R['tcy']}"},
        note="FY2027＝CY2026 下半年（Q 頁）＋ hw × CY2027；FY2028＝（1 − hw）× CY2027 ＋ hw × CY2028")
    row("dcib", "  DCI 基準路徑（未乘情境係數）", "US$B", {"C": "=a25_dci", "D": "=a26_dci", "E": f"=h1dci*(2+{P['h2d']})", "F": f"=E{r}*(1+{P['gd']}_28)", "G": f"=F{r}*(1+{P['gd']}_29)"})
    row("dci", "  DCI（資料中心互連）", "US$B", {c_(y): f"={c_(y)}{R['dcib']}*{c_(y)}{R['uf']}" for y in FYS})
    row("hpcb", "  HPC 基準路徑（未乘情境係數）", "US$B", {"C": "=a25_hpc", "D": "=a26_hpc", "E": f"=h1hpc*(2+{P['h2h']})", "F": f"=E{r}*(1+{P['gh']}_28)", "G": f"=F{r}*(1+{P['gh']}_29)"})
    row("hpc", "  HPC（高效能運算）", "US$B", {c_(y): f"={c_(y)}{R['hpcb']}*{c_(y)}{R['uf']}" for y in FYS})
    row("dc", "資料中心", "US$B", {c_(y): f"={c_(y)}{R['tr']}+{c_(y)}{R['dci']}+{c_(y)}{R['hpc']}" for y in FYS})
    row("ci", "通訊基礎設施", "US$B", {"C": "=a25_ci", "D": "=a26_ci", "E": "=h1ci*(2+h2_ci)", **{c_(y): f"={c_(y - 1)}{r}*(1+g_ci_{y % 100})" for y in (2028, 2029)}})
    row("aio", "車用、工業與其他", "US$B", {"C": "=a25_aio", "D": "=a26_aio", "E": "=h1aio*(2+h2_aio)", **{c_(y): f"={c_(y - 1)}{r}*(1+g_aio_{y % 100})" for y in (2028, 2029)}})
    row("rev", "營收", "US$B", {c_(y): f"={c_(y)}{R['dc']}+{c_(y)}{R['ci']}+{c_(y)}{R['aio']}" for y in FYS})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("dcs", "資料中心 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['dc']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    sec_("C. 損益（non-GAAP）")
    row("gm", "毛利率", "%", {"C": "=a25_gp/a25_rev", "D": "=a26_gp/a26_rev", **{c_(y): f"={P['gm']}_{y % 100}" for y in F}}, fmt="0.0%",
        note="FY2027 起為輸入路徑（成本加成代工；新廠爬坡與料件占比）")
    row("gp", "毛利", "US$B", {c_(y): f"={c_(y)}{R['rev']}*{c_(y)}{R['gm']}" for y in FYS})
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in F}})
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in F}}, fmt="#,##0.000")
    row("op", "營業利益", "US$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("opm", "營業利益率", "%", {c_(y): f"={c_(y)}{R['op']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("oth", "業外淨收益（利息、匯兌、其他）", "US$B", {"C": "=a25_oth", "D": "=a26_oth", **{c_(y): "=int_f" for y in F}}, fmt="#,##0.000")
    row("pre", "稅前利益", "US$B", {c_(y): f"={c_(y)}{R['op']}+{c_(y)}{R['oth']}" for y in FYS})
    row("ni", "淨利", "US$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"={c_(y)}{R['pre']}*(1-({P['tx']}))" for y in F}})
    row("nim", "淨利率", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("sh_", "稀釋股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", "E": "=sh_q1*(1+sh_g)", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in (2028, 2029)}}, fmt="0.0000")
    row("eps", "EPS", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R
# 情境變數別名（A＝高值、C＝低值；Sens 也用）
for nm, base, suf in (("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("PXLO", "pxr", "_lo"), ("GMLO", "gm", "_lo"), ("GMHI", "gm", "_hi"),
                      ("CPLO", "cpo", "_lo"), ("CPHI", "cpo", "_hi")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}_{y}"] = DefinedName(f"{nm}_{y}", attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)
for nm, base, suf in (("GDLO", "g_dci", "_lo"), ("GHLO", "g_hpc", "_lo")):
    for y in (28, 29):
        wb.defined_names[f"{nm}_{y}"] = DefinedName(f"{nm}_{y}", attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Fabrinet 損益（US$B；non-GAAP）。FY2025A、FY2026A（錨點）為實績，FY2027 起為模型。FY(t)＝t−1 年 7 月～t 年 6 月。三情境各一個完整推導區塊（本頁下方），頂端為摘要。顆數係數從 CY2028 開始；FY2027 的情境差異只來自公司特有變數（份額、CPO 份額、毛利率）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數係數 1.0、網通 CPO 情境 A（替代快）；份額、CPO 份額、毛利率取高值", "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85、網通 CPO 情境 B；公司變數取基準", "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6、網通 CPO 情境 C（替代慢）；公司變數取低值"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026A（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.00"), ("g", "  營收年增", "%", "0.0%"), ("dc", "  資料中心", "US$B", "#,##0.00"), ("tr", "    光模組與其他", "US$B", "#,##0.00"),
                                ("dci", "    DCI", "US$B", "#,##0.00"), ("hpc", "    HPC", "US$B", "#,##0.00"), ("ci", "  通訊基礎設施", "US$B", "#,##0.00"),
                                ("aio", "  車用、工業與其他", "US$B", "#,##0.00"), ("gm", "毛利率", "%", "0.0%"), ("opm", "營業利益率", "%", "0.0%"),
                                ("ni", "淨利", "US$B", "#,##0.00"), ("eps", "EPS", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型與實績（校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 B：營收差（US$B）"); PL[f"D{sr}"] = f"=D{PB['B']['rev']}-a26_rev_in"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 B：營業利益差（US$B；公司揭露 0.500063）"); PL[f"D{sr}"] = f"=D{PB['B']['op']}-0.500063"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 B：EPS 差（公司揭露 14.09；四捨五入差應小於 0.01）"); PL[f"D{sr}"] = f"=D{PB['B']['eps']}-a26_eps_in"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 B：CY2026 光模組 − 可插拔市場 × 有效份額"); PL[f"D{sr}"] = f"=D{PB['B']['tcy']}-D{PB['B']['mpl']}*D{PB['B']['sh']}"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 66; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 80
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：市場與份額 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — Fabrinet 在 AI 光模組的有效份額（日曆年）。市場取自 markets/network（加速器顆數 × 每顆光學內容；美元）；份額口徑與中際旭創、Coherent 模型一致（光模組營收 ÷ 可插拔市場），與中際旭創相加須低於 100%"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("mpl", "AI 可插拔光模組市場（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mpl']}", "#,##0.0", "markets/network Mkt 頁"),
        ("mcp", "AI CPO 光學市場（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mcp']}", "#,##0.0", "同上"),
        ("sh", "Fabrinet 光模組有效份額（情境 A）", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "CY2026 以 FY2026 下半年實績推導＋FY2027 Q1 指引推導＋Q2 假設校準；情境 A 份額取高值"),
        ("inl", "中際旭創可插拔有效份額（innolight/dist）", "%", lambda y: f"=inl_{y % 100}", "0.0%", "唯讀上游"),
        ("tot", "兩者合計", "%", lambda y: f"={cyc(y)}{ur - 2}+{cyc(y)}{ur - 1}", "0.0%", "須低於 100%；其餘由新易盛、Coherent、Lumentum 等分配"),
        ("tcy", "Fabrinet 光模組與其他營收（情境 A；日曆年）", "US$B", lambda y: f"=PL!{FC[y]}{A_['tcy']}", "#,##0.00", ""),
        ("oth", "其他廠商＝ 可插拔＋CPO − 中際旭創 − Fabrinet", "US$B", lambda y: f"={cyc(y)}4+{cyc(y)}5-{cyc(y)}7*{cyc(y)}4-{cyc(y)}9", "#,##0.0", "中際旭創以可插拔份額近似（其 CPO 份額不扣）"),
        ("itr", "Fabrinet 資料中心光模組營收 ÷ AI IT 資本支出", "%", lambda y: f"=PL!{FC[y]}{A_['tcy']}/itc_{y % 100}", "0.00%", "markets/network 的 AI IT 資本支出")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
assert UR["sh"] == 6 and UR["inl"] == 7 and UR["tot"] == 8 and UR["tcy"] == 9
UN.cell(ur, 1, "核對：兩者合計最大值 < 100%（1＝通過）"); UN[f"C{ur}"] = f"=IF(MAX(C{UR['tot']}:F{UR['tot']})<1,1,0)"; UR["chk"] = ur; ur += 2
UN.cell(ur, 1, "讀法：Fabrinet 是代工廠，光模組營收含 NVIDIA 等客戶自有品牌的模組，與中際旭創、Coherent 的出貨口徑不同，所以是「有效」份額；NVIDIA 占營收由 27.6% 降到 16.3%，份額相對值與 CPO 份額是 Assumed，是最需要每季檢查的輸入。").font = BOLD
UR["note"] = ur
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 70

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元；non-GAAP EPS）。目標價＝DCF 與 Comparable 加權（第 0 節）；舊法規則值情境加權保留參考；另列「現價 ÷ FY2027、FY2028 EPS（模型）」、市場隱含 EPS 與十年 P/E 規則值。十年 P/E 為 GAAP 年底 P/E（Fabrinet 每年皆獲利，十年皆列入）"
V["A2"] = "循環位置：Fabrinet 不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 Coherent、中際旭創）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(15, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 16, ["項目", "A", "B", "C", "說明"])
V.cell(17, 1, "現價（2026-10-23 收盤）"); V["B17"] = "=px"; V["B17"].number_format = "0.00"
V.cell(18, 1, "現價 ÷ FY2026 EPS（實績）"); V.cell(19, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(20, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(21, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr_, y in ((18, 2026), (19, 2027), (20, 2028), (21, 2029)):
        V[f"{c}{rr_}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr_}"].number_format = "0.0"
V.cell(22, 1, "現價 ÷ 共識未來 12 月 EPS（9/12 × FY2027＋3/12 × FY2028）"); V["B22"] = "=px/(cons_eps27*9/12+cons_eps28*3/12)"; V["B22"].number_format = "0.0"
V.cell(23, 1, "現價 ÷ 52 週高點 − 1"); V["B23"] = "=px/px_hi-1"; V["B23"].number_format = "0%"
V.cell(24, 1, "現價 ÷ FY2027 Q1 EPS 指引中點年化"); V["B24"] = "=px/(q1_eps*4)"; V["B24"].number_format = "0.0"
V["E19"] = "FY2027 三情境差異來自公司特有變數與網通 CPO 情境（顆數係數從 CY2028 開始）"; V["E24"] = "4.175 × 4"
V.cell(26, 1, "2. 循環位置（本模型營收最高的會計年度）").font = BOLD
hdr(V, 27, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 28):
    rr_ = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr_}),{rr_},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(32, 1, "3. 規則值隱含股價（美元；約 FY2027 年底＝2027-06，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 33, ["情境", "FY2027 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 34 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=D{28 + j}").number_format = "0.0"
    V.cell(i, 4, f"=B{i}*C{i}").number_format = "#,##0"
    V.cell(i, 5, f"=D{i}/px-1").number_format = "0%"
    V.cell(i, 6, f"=px/B{i}").number_format = "0.0"
    for c in range(1, 7): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(38, 1, "4. 核心問題：市場隱含的 FY2028 EPS（＝現價 ÷ 假設倍數）與模型的距離").font = BOLD
hdr(V, 39, ["假設倍數（÷ FY2028 EPS）", "市場隱含 FY2028 EPS", "模型 FY2028 EPS（B）", "隱含 ÷ 模型 B − 1", "模型 FY2028 EPS（C）", "隱含 ÷ 模型 C − 1"])
vr = 40; IMPL = {}
for m in (10, 15, 20, "=pe_peak"):
    V.cell(vr, 1, m).number_format = "0.0"
    if not isinstance(m, str): V.cell(vr, 1).font = BLUE
    V.cell(vr, 2, f"=px/A{vr}").number_format = "0.00"
    V.cell(vr, 3, f"=PL!F{SUMR[('B', 'eps')]}").number_format = "0.00"
    V.cell(vr, 4, f"=B{vr}/C{vr}-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[('C', 'eps')]}").number_format = "0.00"
    V.cell(vr, 6, f"=B{vr}/E{vr}-1").number_format = "0%"
    IMPL[m] = vr; vr += 1
V.cell(vr, 1, "讀法：第四列＝十年 P/E 25 分位（見頂倍數）。現價約 14 倍模型 FY2028 EPS（B）；若市場對 FY2028 給見頂倍數，隱含 FY2028 EPS 約等於情境 C（顆數 ×0.6、光模組份額與毛利率低值）——市場定價接近「NVIDIA 光模組份額續降、AI 部署下修」的組合。"); READ_ROW = vr
vr += 2
V.cell(vr, 1, "本版未做先驗：無 Andy 合理 P/E 對照。"); NOPRIOR_ROW = vr
assert vr < 50, vr
# 十年 GAAP 年底 P/E（2016–2025）
PE_CMC = [16.4, 12.0, 18.3, 20.2, 23.0, 25.2, 20.0, 27.2, 24.4, 43.3]
PE_MT = [17.00, 12.21, 18.73, 20.52, 23.37, 25.59, 20.26, 27.50, None, None]
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；GAAP TTM）"},
    {"label": "macrotrends（12 月底 P/E）", "vals": PE_MT, "tag": "Interested-party（資料商；GAAP TTM；2024、2025 年底未列，未找到）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=B{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=B{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=B{VAL['B']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝年底 P/E 依循環位置取分位。兩個資料商來源（GAAP 落後 P/E）；GAAP 倍數套在 non-GAAP EPS 上（non-GAAP 約高 7%），偏誤方向為小幅高估；2025 年 43x 為 AI 光模組題材溢價；前瞻 P/E 歷史待 Bloomberg。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
# 舊法：規則值情境加權（只作參考，不是目標價；建置後與其他內容一起下移）
V.cell(4, 1, "舊法：規則值情境加權（參考）").font = BOLD
hdr(V, 5, ["項目", "數值", "", "", "說明"])
V.cell(6, 1, "規則值情境加權價（參考，非目標價）"); V["B6"] = f"=tp_wA*D{VAL['A']}+tp_wB*D{VAL['B']}+tp_wC*D{VAL['C']}"; V["B6"].number_format = "#,##0"
V["E6"] = "權重 × 各情境隱含價（第 3 節規則值）"
for c in "AB": V[f"{c}6"].fill = KEY
V.cell(7, 1, "規則值情境加權價 ÷ 現價 − 1（參考）"); V["B7"] = "=B6/px-1"; V["B7"].number_format = "0.0%"; V["E7"] = "現價＝In 頁 px（2026-10-23 收盤）"
for rr_, s, lab in ((8, "A", "情境 A 隱含價（順利）"), (9, "B", "情境 B 隱含價（平平）"), (10, "C", "情境 C 隱含價（不如預期）")):
    V.cell(rr_, 1, lab); V[f"B{rr_}"] = f"=D{VAL[s]}"; V[f"B{rr_}"].number_format = "#,##0"
    V[f"E{rr_}"] = f'="FY2027 EPS "&TEXT(B{VAL[s]},"0.00")&" × "&C{28 + SC.index(s)}&" 判定倍數 "&TEXT(C{VAL[s]},"0.0")&"x"'
V.cell(11, 1, "情境權重 A／B／C"); V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0%"
V["E11"] = '="合計 "&TEXT(B11+C11+D11,"0%")&"；In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"'
V.cell(12, 1, "舊法說明")
V["B12"] = ("P/E：各情境 FY2027（2027-06 結束）non-GAAP EPS × 十年（2016–2025 年底）GAAP P/E 依循環位置取分位（營收最高年規則：仍成長＝中段＝中位數；見頂＝25 分位）。"
            "偏誤方向為小幅高估：GAAP 倍數套在較高的 non-GAAP EPS 上；隱含價約為 2027 年中的價格、未折現。")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（以情境 B 為基準；有效份額在每個案例內依 CY2026 重算，FY2026 恆等於實績）：對 FY2027、FY2028 的營收、資料中心營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 資料中心", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 資料中心", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("份額流失（低）", "光模組份額相對值 CY2027 0.80、CY2028 0.70、CY2029 0.60（NVIDIA 份額續降）", dict(rel="RLO")),
         ("份額大增（高）", "份額相對值 1.15／1.25／1.30（新客戶光模組專案、Building 10 產能）", dict(rel="RHI")),
         ("降價快於模型", "單價係數 CY2027 0.90、CY2028 0.80、CY2029 0.72", dict(pxr="PXLO")),
         ("毛利率 低（新廠爬坡）", "毛利率 FY2027 11.8%、FY2028 11.5%、FY2029 11.3%", dict(gm="GMLO")), ("毛利率 高", "毛利率 12.7%／12.8%／12.8%", dict(gm="GMHI")),
         ("CPO 份額 低", "CY2027 0%、CY2028 2%、CY2029 5%", dict(cpo="CPLO")), ("CPO 份額 高", "CY2027 15%、CY2028 20%、CY2029 25%", dict(cpo="CPHI")),
         ("AI 市場 ×0.8", "CY2027 起可插拔光模組市場 ×0.8", dict(mf="mf_lo")), ("AI 市場 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("上半年比重 42%", "日曆年上半年占 42%（FY 認列較晚）", dict(hw="hw_lo")), ("上半年比重 50%", "50%", dict(hw="hw_hi")),
         ("DCI 成長 低", "FY2027 下半年 ÷ 上半年 +0%；FY2028 +5%、FY2029 0%", dict(h2d="h2_dci_lo", gd="GDLO")),
         ("HPC 成長 低（AWS 專案流失）", "FY2027 下半年 ÷ 上半年 −10%；FY2028、FY2029 各 −10%", dict(h2h="h2_hpc_lo", gh="GHLO")),
         ("營業費用率 +0.5pt", "營業費用額外 +0.5% × 營收", dict(opxa="opx_adj_hi")),
         ("計入 Pillar Two 補充稅", "稅率 +9pt（non-GAAP 排除、但為實際稅負）", dict(tx="tx+p2"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.00"), ("D", "dc", 2027, "#,##0.00"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.00"), ("I", "dc", 2028, "#,##0.00"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 B）", "PL 區塊 B", PB["B"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 B）", scn="B", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 A（全部）", "顆數 1.0＋CPO 情境 A＋份額、CPO 份額、毛利率高值", PB["A"], "PL!")
put_case(nn + 1, "情境 C（全部）", "CY2028–29 顆數 ×0.6＋CPO 情境 C＋公司變數低值", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：顆數係數從 CY2028 起作用，FY2028 只含半年 CY2028。FY2027 上半年＝指引推導＋假設，所以份額、市場等案例對 FY2027 只影響下半年；「計入 Pillar Two」顯示 non-GAAP 口徑排除的實際稅負。")
SENS_NOTE = nn + 3
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識、公司指引與產能。只作事後對照，不為對上數字而調參數。共識來自彙整網站（stockanalysis、MarketBeat、tickflow），各家年度與更新日不一"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4; RCR = {}
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True, key=None):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note)
    if key: RCR[key] = rr_
    rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "stockanalysis（2026-10-21）", "", key="rev27")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "stockanalysis（7 位）", "", key="eps27")
rec("FY2027 EPS vs 共識（第二來源）", "=cons_eps27b", pl("eps", 2027), "MarketBeat（2026-10-23）", "", key="eps27b")
rec("FY2028 營收 vs 共識", "=cons_rev28", pl("rev", 2028), "tickflow（3 位；疑似未更新）", "")
rec("FY2028 EPS vs 共識", "=cons_eps28", pl("eps", 2028), "MarketBeat（2026-10-23）", "", key="eps28")
rec("FY2028 EPS vs 共識（第二來源）", "=cons_eps28b", pl("eps", 2028), "tickflow（2 位）", "")
RC.cell(rr_, 1, "B. 公司指引與校準").font = BOLD; rr_ += 1
rec("FY2027 Q1 EPS（模型，營收指引中點、毛利率與費用假設）vs 公司 EPS 指引中點", "=q1_eps", lambda s: "=Q!F19", "公司指引（2026-08-17）", "毛利率、費用、稅率為假設；股數取指引", key="q1")
rec("FY2026 Q3 營收：資料中心＋通訊（回推）＋車用工業（推導）vs 實績", "=q3_rev", lambda s: "=Q!D5+Q!D9+Q!D10", "公司", "恆等（車用工業為差額）；僅確認拆分加總", fmt="#,##0.000")
rec("FY2026 NVIDIA 營收（10-K 16.3%）vs 光模組與其他", "=nv26*a26_rev_in", lambda s: f"=PL!D{PB[s]['tr']}", "10-K", "NVIDIA 營收低於光模組與其他：其餘為其他客戶光模組與資料中心零組件", fmt="#,##0.000")
RC.cell(rr_, 1, "C. 份額、市場與產能").font = BOLD; rr_ += 1
rec("CY2026 Fabrinet＋中際旭創可插拔有效份額", None, lambda s: f"=Units!C{UR['tot']}", "Units 頁", "須低於 100%", fmt="0.0%")
rec("CY2028 Fabrinet＋中際旭創可插拔有效份額", None, lambda s: f"=Units!E{UR['tot']}", "Units 頁（情境 A）", "", fmt="0.0%")
rec("FY2027 資料中心 ÷ CY2027 AI IT 資本支出", None, lambda s: f"=PL!E{PB[s]['dc']}/itc_27", "markets/network", "每 1 美元 AI IT 資本支出分到的 Fabrinet 資料中心營收", fmt="0.00%")
rec("FY2028 營收 vs 產能（現有＋Building 10）", "=cap_b10", pl("rev", 2028), "公司（法說會）", "差距 > 0 表示需要 Building 10 以外的新產能", key="cap28")
rec("FY2029 營收 vs 長期產能目標", "=cap_tot", pl("rev", 2029), "公司（法說會）", "長期目標 125–140 億美元的中點")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以資料中心營收 ÷ 部署 GW 計算每 GW 營收"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.cell(rr_, 1, "FY2027 Q1 季報（預估 2026-11-02 盤後）：公布後以實績取代 Q 頁 Q1 指引推導，並更新資料中心內光模組、DCI、HPC 拆分"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝SEC 8-K 新聞稿與 10-K；法說會內容經第三方轉述，待人工複核）"; SO["A1"].font = BOLD
srcs = [("公司：FY2026 第四季與全年財報新聞稿（8-K Exhibit 99.1；損益、non-GAAP 調節、Pillar Two、Q1 指引、資產負債表）", PR),
        ("公司：FY2026 第三季財報新聞稿（8-K Exhibit 99.1）", PR3),
        ("公司：FY2026 10-K（三類營收比重、客戶集中度、Pillar Two、股數）", TENK),
        ("二手：Webull 法說會逐字稿（FY2026 Q4；類別營收、產能、CPO／NPO、客戶）", WB4),
        ("二手：Nasdaq 法說會重點（類別定義：DCI、HPC 屬資料中心）", NQ4),
        ("二手：Pulse 2.0（Q4 類別營收、DCI 與 HPC 各季）", P2),
        ("二手：Zacks（TradingView；FY2026 Q3 舊分類營收：datacom、telecom、DCI、HPC）", ZK3),
        ("二手：AOL／Motley Fool 法說會逐字稿（FY2026 Q2；HPC、DCI、AWS 第二供應商）", AOL2),
        ("共識：stockanalysis（股價 10-23、forward PE、FY2027 共識、目標價、ROE）", SA + "；" + SAF),
        ("共識：MarketBeat（FY2027／FY2028 EPS、52 週區間、目標價、下次財報估計日 2026-11-02）", MB),
        ("共識：tickflow（FY2028 營收與 EPS）", TF),
        ("資料商：stockanalysis 損益表與資產負債表（FY2022–FY2026；ROE、淨利率）", SAFIN + "；" + SABS),
        ("資料商：companiesmarketcap 年底 P/E", CMC),
        ("資料商：macrotrends 12 月底 P/E、淨利率", MT + "；" + MTM),
        ("同業：stockanalysis forward PE（COHR、LITE、JBL、CLS、FLEX；2026-10-23）", "https://stockanalysis.com/stocks/cohr/statistics/；https://stockanalysis.com/stocks/lite/statistics/；https://stockanalysis.com/stocks/jbl/statistics/；https://stockanalysis.com/stocks/cls/statistics/；https://stockanalysis.com/stocks/flex/statistics/"),
        ("市場層：markets/network、tsmc、innolight 成品（本 repo，唯讀）；markets/memory SCEN_UNITS（核對）", ""),
        ("待補（人工）：光模組（datacom）新分類金額；NVIDIA 光模組營收；CPO 營收；泰國 BOI 稅率明細；共識家數與 FY2028 營收第二來源；前瞻 P/E 歷史（Bloomberg）；FY2027 Q1 財報日期公司公告", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Fabrinet（FN，NYSE）收支與估值模型 v0.1.2（AI 半導體 Project：網通次產業內複製——光模組與光元件代工，樣板 Coherent）　2026-10-26", True),
 ("版本：v0.1.1（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）；v0.1（2026-10-25 初版；三情境延伸到 FY2032、目標價＝DCF 與 Comparable 加權）。Fabrinet 的光模組有效份額口徑與中際旭創、Coherent 模型一致（光模組營收 ÷ 可插拔市場），Units 頁檢查與中際旭創相加低於 100%。", False),
 ("模型命題：Fabrinet 的資料中心營收＝AI 可插拔光模組市場（本專案網通模型）× Fabrinet 光模組有效份額（NVIDIA 等客戶代工）× 單價係數＋CPO 光學 × CPO 份額＋DCI＋HPC（乘同字母 AI 情境係數）；通訊基礎設施、車用工業另依成長率推。股價自 52 週高點 749 跌到約 315（−58%），只約 17 倍共識未來 12 月 EPS：市場在問 NVIDIA 光模組份額流失（占營收 27.6% → 16.3%）後，DCI、HPC 與新客戶能否接上成長。", False),
 ("核心問題：NVIDIA 份額下滑而 DCI、HPC 高成長——市場隱含的 FY2028 EPS 要多低，光模組份額、AI 市場、毛利率三條路徑各要不利到什麼程度？", False), ("", False),
 ("會計年度：FY(t)＝t−1 年 7 月～t 年 6 月（6 月最後一個星期五結束）；FY2026＝2025-06～2026-06-26（實績）、FY2027＝2026-06～2027-06。日曆年市場以「上半年比重 hw」換算：FY(t)＝（1 − hw）× CY(t−1)＋hw × CY(t)，FY2027 的 CY2026 下半年用 Q1 指引推導＋Q2 假設。口徑：non-GAAP（排除股票薪酬、Pillar Two 補充稅、非上市股權重估等）。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("光模組與其他：markets/network 的 AI 可插拔光模組市場 × 同字母顆數係數 × Fabrinet 有效份額（CY2026 校準）× 份額相對值（Assumed）× 單價係數（基準 1）＋ CPO 光學市場 × CPO 份額（Assumed）。DCI、HPC：FY2027 上半年推導 × 下半年成長，FY2028 起年增，乘同字母 AI 情境係數。", False),
 ("通訊基礎設施（電信系統、衛星通訊、電信元件）、車用工業與其他（EV 充電、LiDAR、工業雷射）：FY2027 上半年 × 下半年成長，FY2028 起年增（Assumed；各情境相同）。毛利率：FY2027 起輸入路徑（情境 A 高、C 低）。營業費用：彈性 phi（Derived）。稅率 5%（不含 Pillar Two）。", False),
 ("校準：FY2026 全年實績；資料中心內光模組、DCI、HPC 由法說會金額推導（Q 頁）。不用共識（只在 Recon 對帳）。", False),
 ("估值：目標價（Val 第 0 節）＝0.5 × DCF（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，情境權重 0.3／0.5／0.2；情境價值下限 0）＋0.5 × Comparable（市場主流同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）。舊法規則值情境加權保留參考（不是目標價）；現價 ÷ FY2027、FY2028 EPS；市場隱含 FY2028 EPS（Val 第 4 節）；十年 GAAP P/E 規則值（兩個資料商）× FY2027 EPS。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。", False),
 ("", False),
 ("已知限制：(1) 光模組份額相對值、CPO 份額、毛利率路徑皆為 Assumed，是 FY2028 EPS 最敏感的輸入（Sens）；(2) 資料中心內光模組、DCI、HPC 金額為法說會轉述與推導，FY2027 Q1 季報（預估 2026-11-02）後更新；(3) 代工營收含料件轉嫁，有效份額與模組廠出貨份額口徑不同；(4) non-GAAP 排除 Pillar Two 補充稅（約每股 1.6 美元／年），DCF 偏高估；(5) Building 10 等資本支出使 FY2026 自由現金流近 0，DCF 以 g／ROE 再投資率近似，未建資本支出；(6) R4 產業部署路徑未建（Recon 預留列）。", False),
 ("情境對應：A＝順利（顆數 1.0、份額 1.15／1.25／1.30、CPO 15／20／25%、毛利率 12.7–12.8%）；B＝平平（顆數 0.8／0.85、份額 1.00／1.00／0.95、CPO 5／10／15%、毛利率 12.3%）；C＝不如預期（顆數 0.6、份額 0.80／0.70／0.60、CPO 0／2／5%、毛利率 11.3–11.8%）。FY2030 起為收斂法，非驅動因子推導（PL 灰底）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170


# ================= FY2030–FY2032 收斂法；目標價＝DCF 與 Comparable 加權 =================
# 工作單：Andy 2026-10-25 11:18／11:20（三情境延伸到 FY2032、收斂法）、11:28（目標價＝DCF＋Comparable）、11:53（tp_w_dcf 0.5；A 股與其他市場同業分開）、12:22（DCF 下限 0）
import re as _re
from copy import copy as _copy
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
FCX = dict(FC); FCX.update({2030: "H", 2031: "I", 2032: "J"})
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
INR = {IN.cell(i, 1).value: i for i in range(4, IN.max_row + 1) if IN.cell(i, 1).value}
for _k, _i in INR.items():
    if str(_k).startswith(("tp_", "fade_", "roe_")) or _k in ("cons_tp", "cons_feps", "nim10"):
        for _j in range(1, 10): IN.cell(_i, _j).fill = KEY

def _move_notes(ws):
    """說明欄（第 8 欄，H）移到第 11 欄（K），讓出 H–J；表頭列補上收斂年度。"""
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 8)
        if c.value is None: continue
        assert not (isinstance(c.value, str) and c.value.startswith("=")), (ws.title, r, c.value)
        k = ws.cell(r, 11); k.value = c.value; k.font = _copy(c.font); k.fill = _copy(c.fill)
        c.value = None; c.font = Font(); c.fill = PatternFill()
        if str(ws.cell(r, 1).value or "") in ("項目", "情境／項目"):
            for j, y in zip((8, 9, 10), FADE):
                h = ws.cell(r, j, f"FY{y}（收斂）"); h.font = BOLD; h.fill = HDR
    ws.column_dimensions["K"].width = ws.column_dimensions["H"].width or 70
    for col in "HIJ": ws.column_dimensions[col].width = 14

def _fade(ws, R, s, eps_f):
    """FY2030–32：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到）；非驅動因子推導。"""
    for y in FADE:
        c, pc, k = FCX[y], FCX[y - 1], y - 2029
        fs = {"g": f"=$G{R['g']}+(fade_g{s}-$G{R['g']})*{k}/3", "rev": f"={pc}{R['rev']}*(1+{c}{R['g']})",
              "nim": f"=$G{R['nim']}+(fade_m{s}-$G{R['nim']})*{k}/3", "ni": f"={c}{R['rev']}*{c}{R['nim']}",
              "sh_": f"={pc}{R['sh_']}*(1+sh_g)", "eps": "=" + eps_f(c, R), "epsg": f"={c}{R['eps']}/{pc}{R['eps']}-1"}
        for key, f in fs.items():
            ws[f"{c}{R[key]}"] = f; ws[f"{c}{R[key]}"].number_format = ws[f"G{R[key]}"].number_format; ws[f"{c}{R[key]}"].fill = FADEFILL

def _shift_rows(wb, sheet, start, n):
    """把 sheet 第 start 列起的內容整體下移 n 列，並同步改寫所有公式與定義名稱中指向這些列的參照（等同插入列）。"""
    ws = wb[sheet]
    cell_ = r"\$?[A-Z]{1,3}\$?\d+"
    pat = _re.compile(r"(?<![A-Za-z0-9_.!$:])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(" + cell_ + r")(?::(" + cell_ + r"))?(?![A-Za-z0-9_(!])")
    one = _re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")
    def fix(f, own):
        parts = f.split('"')
        for i in range(0, len(parts), 2):
            def rep(m):
                pre = m.group(1)
                tgt = (pre is None and own) or (pre is not None and pre.strip("'!") == sheet)
                def sh(x):
                    a, rr = one.fullmatch(x).groups()
                    rr = int(rr)
                    return f"{a}{rr + n if (tgt and rr >= start) else rr}"
                return (pre or "") + sh(m.group(2)) + (":" + sh(m.group(3)) if m.group(3) else "")
            parts[i] = pat.sub(rep, parts[i])
        return '"'.join(parts)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = fix(c.value, w.title == sheet)
    for nm in list(wb.defined_names.keys()):
        d = wb.defined_names[nm]
        if sheet + "!" in d.attr_text: d.attr_text = fix(d.attr_text, False)
    cells = [(c.row, c.column) for row in ws.iter_rows(min_row=start) for c in row if c.value is not None or c.has_style]
    for r, col in sorted(cells, reverse=True):
        src = ws.cell(r, col); dst = ws.cell(r + n, col)
        dst.value = src.value; dst._style = _copy(src._style); dst.number_format = src.number_format
        src.value = None; src._style = _copy(ws.cell(1, 30)._style)

def _tp_section(V, eps_ref, fwd, fwd_note, peers, peer_note, cons_note, method, pe_rng, fmt, cons_feps_note):
    """Val 第 0 節（第 4–48 列）：0.1 DCF、0.2 Comparable、0.3 結論。列標籤依工作單逐字；B／C／D＝情境 A 順利／B 平平／C 不如預期。"""
    T = {}
    def tset(r, key, label, fs, nf, note=""):
        V.cell(r, 1, label); T[key] = r
        for c, f in zip("BCD", fs):
            if f is None: continue
            V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = nf
        if note: V[f"E{r}"] = note
    V.cell(4, 1, "0. 目標價").font = BOLD
    V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
    hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
    V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法（PL 頁灰底）"
    for i, y in enumerate(YRS6):
        tset(8 + i, f"eps{y}", f"  FY{y}", [f"={eps_ref(s, y)}" for s in SC], "0.00")
    V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）")
    tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=IF(tp_roe<=tp_g,0.8,MIN(tp_g/tp_roe,0.8))", None, None], "0.0%", "ROE＝In 頁 tp_roe；tp_roe ≤ tp_g 時取上限 0.8")
    for i, y in enumerate(YRS6):
        tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{8 + i}*(1-$B$15)" for c in "BCD"], "0.00")
    tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], fmt,
         "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027；小於 0 時取 0（股東有限責任）")
    tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], fmt)
    tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], fmt, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
    V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
    hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
    assert 3 <= len(peers) <= 5, peers
    for i, (nm, tk, pe, dt, src) in enumerate(peers, 28):
        V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
    tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{27 + len(peers)})", None, None], "0.0", peer_note)
    tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）",
         ["=" + "+".join(f"tp_w{s}*(" + "+".join(f"{w}*{c}{8 + YRS6.index(y)}" for y, w in fwd) + ")" for s, c in zip(SC, "BCD")), None, None], "0.00", fwd_note)
    tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], fmt, "同業中位數 × 前瞻 EPS")
    V.cell(37, 1, "0.3 結論").font = BOLD
    hdr(V, 38, ["項目", "數值", "", "", "說明"])
    tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], fmt, "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
    tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], fmt)
    tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], fmt, cons_note)
    tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px")
    tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
    tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", cons_feps_note)
    tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
    V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
    tset(46, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D23,"{fmt}")&"～"&TEXT(B23,"{fmt}")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
    tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
         [f"=IF(px/{c}9>MAX({pe_rng}),1,IF(px/{c}9<MIN({pe_rng}),0,PERCENTRANK({pe_rng},px/{c}9)))" for c in "BCD"], "0%",
         "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 6 節採用來源）的分位；高於十年最大值＝100%；只作檢查，不進目標價")
    V.cell(48, 1, "目標價方法"); V["B48"] = method
    for r_ in (24, 35, 39):
        V[f"B{r_}"].fill = KEY
    return T

def _tp_sens(SE, T):
    """Sens：目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.3／0.7）；放在案例摘要與第 40 列推導區塊之間。"""
    last = max(r for r in range(1, 40) if any(SE.cell(r, j).value is not None for j in range(1, 13)))
    t0 = last + 2
    SE.cell(t0, 1, "目標價參數（只改 DCF 參數與權重；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
    hdr(SE, t0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
    def dcf_f(c, r, g, reinv):
        cf = lambda i: f"Val!{c}{8 + i}*(1-{reinv})"
        pv = "+".join(f"{cf(i)}/(1+{r})^{i + 1}" for i in range(6))
        return f"MAX(0,{pv}+{cf(5)}*(1+{g})/({r}-{g})/(1+{r})^6)"
    cases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
             ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
             ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
    for k, (nm, st, kw) in enumerate(cases):
        i = t0 + 2 + k; r = kw.get("r", "tp_r"); g = kw.get("g", "tp_g"); w = kw.get("w", "tp_w_dcf")
        reinv = f"IF(tp_roe<={g},0.8,MIN({g}/tp_roe,0.8))"
        SE.cell(i, 1, nm); SE.cell(i, 2, st)
        SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_f(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
        SE[f"D{i}"] = f"=Val!B{T['comptp']}"
        SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
        SE[f"F{i}"] = f"=E{i}/E${t0 + 2}-1"
        for c in "CDE": SE[f"{c}{i}"].number_format = "0.00"
        SE[f"F{i}"].number_format = "0.0%"
    assert t0 + 1 + len(cases) < 40, t0
    return t0, t0 + 1 + len(cases)

def _fade_all(eps_f, nim10):
    """PL：三情境推導區塊與頂端摘要延伸到 FY2032；In 頁 fade_m 連到 PL 頁 FY2029 淨利率與十年平均。"""
    _move_notes(PL)
    for s in SC:
        _fade(PL, PB[s], s, eps_f)
        for key in ("rev", "g", "nim", "ni", "eps", "epsg"):
            if (s, key) not in SUMR: continue
            for y in FADE:
                cc = PL[f"{FCX[y]}{SUMR[(s, key)]}"]; cc.value = f"={FCX[y]}{PB[s][key]}"
                cc.number_format = PL[f"G{SUMR[(s, key)]}"].number_format; cc.fill = FADEFILL
    nimr = {s: f"PL!$G${PB[s]['nim']}" for s in SC}
    IN.cell(INR["fade_mA"], 5, f"={nimr['A']}")
    IN.cell(INR["fade_mB"], 5, f"=({nimr['B']}+{nim10})/2")
    IN.cell(INR["fade_mC"], 5, f"={nim10}")
    PL["A2"] = ("灰底＝FY2030–FY2032 收斂法，非驅動因子推導（上游市場層只到 2029）：營收年增由 FY2029 線性收斂到 fade_g（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt），"
                "淨利率由 FY2029 線性收斂到 fade_m（A＝FY2029、B＝FY2029 與十年平均的中點、C＝十年平均；In 頁 K 節），FY2032 達到；股數依 sh_g")

def _rule_pe_range(V, ncols=1):
    """第 6 節十年倍數表（2016–2025 列）中採用來源的欄位範圍（PERCENTRANK 用；ncols＞1 時各來源合併成一個分布）。"""
    yr = [r for r in range(1, V.max_row + 1) if V.cell(r, 1).value in [str(y) for y in range(2016, 2026)]]
    yr = [r for r in yr if r > 50][:10]
    assert len(yr) == 10 and yr[-1] - yr[0] == 9, yr
    return f"$B${yr[0]}:${chr(65 + ncols)}${yr[-1]}"

_fade_all(lambda c, R: f"{c}{R['ni']}/{c}{R['sh_']}", "nim10")
_shift_rows(wb, "Val", 4, 47)
PEERS = [('Coherent', 'COHR', 22.30, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'),
         ('Lumentum', 'LITE', 34.71, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'),
         ('Jabil', 'JBL', 15.49, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'),
         ('Celestica', 'CLS', 16.65, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）'),
         ('Flex', 'FLEX', 14.74, '2026-10-23', 'stockanalysis forward PE（分析師共識；價格日期如左）')]
TPT = _tp_section(V, lambda s_, y: f"PL!{FCX[y]}{SUMR[(s_, 'eps')]}", [(2028, '8/12'), (2029, '4/12')],
    '約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 8／12 × FY2028＋4／12 × FY2029（會計年度 6 月底結束）',
    PEERS,
    '美股光通訊（Coherent、Lumentum）與電子代工（Jabil、Celestica、Flex）同業；不放 A 股（中際旭創、新易盛另見各自模型）',
    'stockanalysis：9 位分析師平均 687.78（2026-10-21 更新；MarketBeat 10 位平均 650.63）',
    'DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年 GAAP 平均（再投資率約兩成），FY2027–32 明確期＋永續，FY2030 起為收斂法，情境價值下限 0。Comparable：Coherent、Lumentum、Jabil、Celestica、Flex 前瞻 P/E 中位數 × 12 個月後前瞻 EPS（non-GAAP）。偏誤：會計年度 6 月底結束，DCF 的 t＝1 對 FY2027，比基準日早半年（偏高估）；non-GAAP 排除 Pillar Two 補充稅（偏高估約 9%）；資本支出高於折舊、未建模（偏高估）。情境對應：A＝順利、B＝平平、C＝不如預期。',
    _rule_pe_range(V, 2), '#,##0.00',
    "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝In 頁 cons_feps")
TPS0, TPS_LAST = _tp_sens(SE, TPT)

# ---------------- v0.1.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 14:00、14:17；既有數值與儲存格位置不變） ----------------
_UPW = [r_ for r_ in range(1, 60) if V.cell(r_, 1).value == "上檔空間：加權目標價"]
assert _UPW == [42], _UPW
GR_ROW, RC_ROW = 49, 50   # 0.3 結論後的空白預留列（不移動其他列）
for _r in (GR_ROW, RC_ROW):
    assert all(V.cell(_r, _c).value is None for _c in range(1, V.max_column + 1)), _r
V.cell(GR_ROW, 1, "上檔空間級距")
V.cell(GR_ROW, 2, '=IF(B42>0.5,"大於 50%",IF(B42>0.25,"25%–50%",IF(B42>=0,"0%–25%",IF(B42>=-0.25,"0%～−25%",IF(B42>=-0.5,"−25%～−50%","小於 −50%")))))').number_format = "@"
V.cell(GR_ROW, 5, "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
V.cell(RC_ROW, 1, "投資建議")
V.cell(RC_ROW, 2, '=IF(B42>0.5,"強力買進",IF(B42>0.25,"買進",IF(B42>=-0.25,"中立",IF(B42>=-0.5,"放空","強力放空"))))').number_format = "@"
V.cell(RC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = [
       "目標價 US$373（DCF 350、Comparable 396 各半），比現價 314.94（2026-10-23）高 18.5%，級距 0%–25%；Comparable 單獨為 +25.9%。本模型 12 個月後前瞻 EPS 23.81 比共識前瞻 EPS 21.27 高 11.9%，但同業（Coherent、Lumentum、Jabil、Celestica、Flex）前瞻 P/E 中位數 16.7 倍，遠低於共識目標價 687.78 隱含的 32.3 倍（高 94.2%）；共識目標價比現價高 118.4%，與本模型的差距幾乎全在倍數。",
       "情境分歧大：情境 A／C（含網通 CPO 情境、份額、CPO 份額、毛利率）FY2028 EPS 比 B 高 26.2%／低 20.7%，DCF 價格帶 164～564；單一變數中計入 Pillar Two 補充稅 −9.5%、光模組份額流失 −7.7%、毛利率低（新廠爬坡）−6.9%。",
       "追蹤 FY2027 Q1 財報（預估 2026-11-02）：資料中心內光模組、DCI、HPC 金額為推導值，NVIDIA 光模組營收未揭露。網通 CPO 已改三情境（v0.1.2；A＝替代快、C＝慢）：Fabrinet 同時代工可插拔與 CPO（情境 A 取 CPO 份額高值），CPO 快淨為加分——情境 A FY2028 EPS 由 26.89 升到 27.81（+3.4%）、C 由 17.44 升到 17.48（+0.2%）、B 不變；CPO 份額低／高時 FY2028 EPS −2.0%／+3.0%。"
      ]   # 重點觀察（與 README.md 相同文字）
_n0 = RD.max_row
RD.cell(_n0 + 1, 1, "v0.1.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（Val 頁 0.3 結論加「上檔空間級距」「投資建議」兩列、README 頁加重點觀察三點；數值不變，既有儲存格位置不變）。")
RD.cell(_n0 + 2, 1, f'="投資建議："&Val!B{RC_ROW}&"（依上檔空間固定規則）"').font = BOLD
for _i, _t in enumerate(OBS, 1): RD.cell(_n0 + 2 + _i, 1, f"重點觀察 {_i}：" + _t)


# ---------------- 分部原則第 3 條（2026-10-28）：核對列與宣告（追加於工作表最後，不位移既有列，數值不變） ----------------
import os as _os2, sys as _sys2
_sys2.path.insert(0, _os2.path.join(_os2.path.dirname(_os2.path.dirname(_os2.path.abspath(__file__))), "tools", "segment_check"))
from addrow import add_check
add_check(Q, seg_rows=[5, 9, 10], total_row=4, cols="BCDEFG", rnd=3, note="分部原則第 3 條（2026-10-28 追加）：公司揭露的終端市場營收類別（單一應報導部門）；分部宣告見 segments.json")
RD.cell(RD.max_row + 1, 1, "v0.1.3（2026-10-28）：分部原則重建（數值不變）。Q 頁最後追加核對列「檢查：分部合計 − 合併營收」；官方分部與子列宣告見 segments.json。")
if isinstance(RD.cell(1, 1).value, str): RD.cell(1, 1).value = RD.cell(1, 1).value.replace("v0.1.2", "v0.1.3")

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# ---------------- v0.1.2（2026-10-26）：網通 CPO 三情境（同字母）；In 頁在現有最後一列之後新增，不移動既有列 ----------------
_ri = IN.max_row + 2
IN.cell(_ri, 1, "網通 CPO 三情境（v0.1.2；markets/network v0.2.3：A＝替代快、B＝平平、C＝慢；情境 A 為上方 A 節，B、C 為下列，各再乘同字母情境顆數係數）").font = BOLD; _ri += 1
for _kb, _lab, _kind in [('mpl', 'AI 可插拔光模組市場', 'mpl'), ('mcp', 'AI CPO 光學市場', 'mcp')]:
    for _s in ("B", "C"):
        for _j, _y in enumerate((26, 27, 28, 29)):
            _k = f"{_kb}{_s}_{_y}"
            IN.cell(_ri, 1, _k); IN.cell(_ri, 2, f"情境 {_s}：{_lab} CY20{_y}"); IN.cell(_ri, 3, "US$B")
            IN.cell(_ri, 5, NETS[(_kind, _s)][0][_j]).font = BLUE
            IN.cell(_ri, 7, "Derived"); IN.cell(_ri, 8, NETS[(_kind, _s)][1])
            IN.cell(_ri, 9, f"網通模型 CPO 情境 {_s}（技術替代屬行業情境，公司讀同字母）" if _y == 26 else "同上")
            wb.defined_names[_k] = DefinedName(_k, attr_text=f"In!$E${_ri}"); _ri += 1
_n1 = RD.max_row
RD.cell(_n1 + 1, 1, 'v0.1.2（2026-10-26）：網通模型 v0.2.3 的 CPO 滲透改三情境（A＝替代快、B＝平平、C＝慢），本模型 A／B／C 改讀同字母的可插拔與 CPO 市場（In 頁末新增 B、C 兩組；情境 A 沿用 A 節 mpl_／mcp_，數值改為網通情境 A）；FY2027 起三情境不再相同。其他輸入不變。')
# 相對倍數調整（Andy 2026-10-27「照建議做」；tools/relative_multiple.py，2026-10-29 接入）
import os, sys  # noqa: E401
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools"))
import relative_multiple as RMX  # noqa: E402
RMX.apply(wb, "fabrinet")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tools", "in_layout")); import apply_in_layout  # noqa: E402
apply_in_layout.apply(wb, "v0.1.4", "2026-10-30", split=True, split_skip=[])  # v0.1.4：In 頁版面；v0.1.5（2026-10-31）：第 3 點情境專屬列拆分（數值不變）
wb.save(sys.argv[1])
if os.environ.get("FN_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK_ROW, "SENSN": SENS_N, "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE,
               "VAL": VAL, "IMPL": {str(k): v for k, v in IMPL.items()}, "READ": READ_ROW, "NOPRIOR": NOPRIOR_ROW, "RCR": RCR, "RC_LAST": RC_LAST, "INROW": INROW, "RULE": rule_rows,
               "TPT": TPT, "TPS0": TPS0, "TPS_LAST": TPS_LAST},
              open(os.environ["FN_ROWMAP"], "w"), ensure_ascii=False, indent=1)
