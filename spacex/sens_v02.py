# 敏感度：一次改一個 In 頁輸入（敏感度低／高值），整本以 LibreOffice 重算，讀 Val 頁加權目標價，寫成 JSON 供 build_spacex.py --sens 使用。
# 用法：python3 spacex/sens_v02.py <工作資料夾> <輸出 sens.json>（環境變數 JOBS＝平行數，預設 2）
import json, os, subprocess, sys, shutil, tempfile
from concurrent.futures import ThreadPoolExecutor
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WD, OUT = sys.argv[1], sys.argv[2]
os.makedirs(WD, exist_ok=True)
BUILD = os.path.join(HERE, "build_spacex.py")


def build(name, sets):
    raw = os.path.join(WD, f"{name}.xlsx")
    cmd = [sys.executable, BUILD, raw] + sum((["--set", f"{k}={v}"] for k, v in sets), [])
    env = dict(os.environ); env["ROWMAP"] = os.path.join(WD, f"{name}.rowmap.json")
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, env=env)
    return raw


def recalc(raw):
    out = os.path.join(WD, "out"); os.makedirs(out, exist_ok=True)
    prof = tempfile.mkdtemp(prefix="lo_spcx_", dir=WD)
    env = dict(os.environ); env["LC_ALL"] = "C.UTF-8"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", out, raw], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=env, timeout=900)
    shutil.rmtree(prof, ignore_errors=True)
    return os.path.join(out, os.path.basename(raw))


def tp_of(xlsx, rowmap):
    rm = json.load(open(rowmap, encoding="utf-8"))
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    return {"tp": wb["Val"].cell(rm["VR"]["tp"], 2).value, "gateB": wb["CF"].cell(rm["GATE"], 3).value,
            "s2B": wb["Seg"][f"I{rm['SR']['B']['s2']}"].value, "plug27B": wb["CF"][f"H{rm['CR']['B']['plug']}"].value,
            "rec": wb["Val"].cell(rm["VR"]["rec"], 2).value}


def run(name, sets):
    raw = build(name, sets)
    return tp_of(recalc(raw), os.path.join(WD, f"{name}.rowmap.json"))


base_raw = build("base", [])
rm = json.load(open(os.path.join(WD, "base.rowmap.json"), encoding="utf-8"))
IN = openpyxl.load_workbook(base_raw)["In"]
row = rm["IN_ROW"]
jobs = {"base|base": []}
for key, side, val, _r in rm["SE_ROWS"]:
    if f"{key}|{side}" in rm["SE_SPECIAL"]:
        jobs[f"{key}|{side}"] = [tuple(x) for x in rm["SE_SPECIAL"][f"{key}|{side}"]]
        continue
    v = IN.cell(row[key], 4 if val == "lo" else 6).value if val in ("lo", "hi") else val
    jobs[f"{key}|{side}"] = [(key, v)]
res = {}
with ThreadPoolExecutor(max_workers=int(os.environ.get("JOBS", "2"))) as ex:
    futs = {k: ex.submit(run, f"c{i:02d}", sets) for i, (k, sets) in enumerate(jobs.items())}
    for k, f in futs.items():
        o = f.result()
        res[k] = {**o, "set": jobs[k]}
        print(k, jobs[k], round(o["tp"], 2), round(o["gateB"], 1), round(o["s2B"], 3), flush=True)
json.dump(res, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
