# SpaceX 個股模型 v0.2 成品核對：公式錯誤 0、Q／Seg／PL／CF／Val 核對列 0（含分部對帳、子列合計、命題子列合計）、敏感度未過期、固定列標籤存在。
# 用法：python3 spacex/check_v02.py <重算後的 xlsx>；失敗時非零結束。
import sys
import openpyxl

X = sys.argv[1]
wb = openpyxl.load_workbook(X, data_only=True)
fail = []
ERR = ("#REF", "#DIV", "#VALUE", "#NAME", "#N/A", "#NUM", "Err:")
n_err = 0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(ERR):
                n_err += 1
                if n_err <= 10: fail.append(f"公式錯誤 {ws.title}!{c.coordinate} {c.value}")
print(f"公式錯誤：{n_err} 格")


def rows_with(ws, prefix):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith(prefix):
            yield r


def nums(ws, r):
    return [c.value for c in ws[r][1:] if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


n_chk = 0
seg_rec = 0
for sh in ("Q", "Seg", "PL", "CF", "Val"):
    ws = wb[sh]
    for r in rows_with(ws, "檢查："):
        lab = str(ws.cell(r, 1).value)
        if "應為 1" in lab: continue
        vs = nums(ws, r)
        n_chk += 1
        if any(("分部合計" in lab or "子列合計" in lab) for _ in [0]): seg_rec += 1
        if any(abs(v) > 1e-6 for v in vs):
            fail.append(f"{sh}!A{r}「{lab}」有非零值 {[round(v, 6) for v in vs if abs(v) > 1e-6][:5]}")
print(f"核對列：{n_chk} 列（Q、Seg、PL、CF、Val；含三情境）；其中分部對帳與子列合計 {seg_rec} 列")
V = wb["Val"]
need = ["目標價（DCF 與 Comparable 加權）", "上檔空間：加權目標價", "投資建議", "上檔空間級距", "DCF 目標價（情境加權）", "Comparable 目標價（市場主流同業法）",
        "價格帶（DCF 不如預期～順利）", "現價（2026-10-29 收盤）", "市場共識目標價"]
labs = {V.cell(r, 1).value: r for r in range(1, V.max_row + 1)}
for n in need:
    if n not in labs: fail.append(f"Val 缺列「{n}」")
cf = wb["CF"]
if "收支關卡：自由現金流（收支關卡口徑，五年合計）" not in {cf.cell(r, 1).value for r in range(1, cf.max_row + 1)}:
    fail.append("CF 缺「收支關卡：自由現金流（收支關卡口徑，五年合計）」")
se = wb["Sens"]
for r in range(1, se.max_row + 1):
    if str(se.cell(r, 1).value or "").startswith("敏感度重算時的基準值"):
        d = se.cell(r, 5).value
        if not isinstance(d, (int, float)) or abs(d) > 1e-6:
            fail.append(f"Sens 敏感度已過期或缺：基準差 {d}")
# 二修：pace 的敏感度低／高（建置時以 Python 重算 Space＋Connectivity Q4 資本支出）須與重算後模型一致
IN = wb["In"]; K = {IN.cell(r, 1).value: r for r in range(4, IN.max_row + 1) if IN.cell(r, 1).value}
TH = wb["Thesis"]
q4c = next((TH.cell(r, 6).value for r in range(1, TH.max_row + 1) if str(TH.cell(r, 1).value or "").startswith("Q4 26 SI 現金資本支出")), None)
cg = next((TH.cell(r, 3).value for r in range(1, TH.max_row + 1) if str(TH.cell(r, 1).value or "").startswith("每 GW 投資（2026 新增")), None)
if "si_sh_cx" not in K and q4c is not None:
    E_ = lambda k, c=5: IN.cell(K[k], c).value
    x = E_("cx_q4") - q4c
    for c, cx in ((4, E_("cx_q4", 4)), (6, E_("cx_q4", 6))):
        want = 4 * (cx - x + E_("ppefin_9m") / 3) / cg
        if abs(E_("pace", c) - want) > 5e-4: fail.append(f"In pace 敏感度欄 {c} {E_('pace', c)} ≠ 模型同一拆法 {want:.4f}")
    if abs(E_("si_cx_rr") - (q4c + E_("ppefin_9m") / 3)) > 1e-9: fail.append("si_cx_rr ≠ Q4 SI 現金資本支出＋以融資取得 PP&E 季均")
    print(f"pace 一致性：Space＋Connectivity Q4 {x:.4f}；pace 低／基準／高 {E_('pace', 4):.4f}／{E_('pace'):.4f}／{E_('pace', 6):.4f}")
tp = V.cell(labs["目標價（DCF 與 Comparable 加權）"], 2).value
up = V.cell(labs["上檔空間：加權目標價"], 2).value
rec = V.cell(labs["投資建議"], 2).value
print(f"加權目標價 {tp:.2f}；上檔空間 {up:+.1%}；投資建議 {rec}")
if fail:
    print("失敗：")
    for f in fail: print("  -", f)
    sys.exit(1)
print("全部通過")
