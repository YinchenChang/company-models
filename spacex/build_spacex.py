# SpaceX（Nasdaq: SPCX，含 xAI）個股模型 v0.1：命題模組＋官方營收分部 → 損益表（docs/plan/個股模型原則_命題加營收分部.md）
# 工作單：spacex/docs/workorders/20261030_SpaceX_v0.1_個股模型.md（B 段；含「A 段審查結果（chat 端 2026-10-30）」）
# 架構範本：alphabet/build_alphabet.py（Alphabet v0.3.2）；SI 算力命題參考 OpenAI v0.6.4（MW 錨定）與 CoreWeave（算力出租）。
# 用法：python3 spacex/build_spacex.py raw.xlsx [--sens sens.json] [--set 鍵=值 …]；再以 LibreOffice headless 重算（spacex/make_v01.sh 一次跑完）。
# 金額：十億美元（US$B）；每股：美元；股數：十億股；GW＝IT 口徑名目算力（nameplate compute，不乘 PUE）。日曆年＝會計年度。
# 欄位（模型頁共用）：C FY2025*（二手轉述 S-1）、D 9M25A、E 9M26A（內部鍵 "1H"）、F Q4 26E（"2H"）、G FY2026E（＝E＋F）、H–M FY2027–FY2032。評價日 2026-09-30。
import sys, os, json, argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ap = argparse.ArgumentParser()
ap.add_argument("out")
ap.add_argument("--sens", help="敏感度重算結果（make_v01.sh 產生）；省略時 Sens 頁結果欄留空")
ap.add_argument("--set", action="append", default=[], help="覆寫 In 頁輸入：鍵=值（敏感度重算用）")
A = ap.parse_args()
OVR = {}
for s in A.set:
    k, v = s.split("=", 1); OVR[k] = float(v)

BLUE = Font(color="0000FF"); BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7"); KEY = PatternFill("solid", fgColor="FFF2CC"); SUB = PatternFill("solid", fgColor="E2EFDA")
wb = openpyxl.Workbook()

FD = json.load(open(os.path.join(HERE, "data", "spacex_facts_2026q3.json"), encoding="utf-8"))
F = {it["id"]: it for it in FD["items"]}
TKV = "v21.7"
TK = json.load(open(os.path.join(HERE, "data", f"tokenomics_snapshot_{TKV}.json"), encoding="utf-8"))
GEN = {x["code"]: x["name"] for x in TK["source"]["generations"]}
PSI = json.load(open(os.path.join(HERE, "data", "peers_si_20261030.json"), encoding="utf-8"))
PCN = json.load(open(os.path.join(HERE, "data", "peers_connectivity_20261030.json"), encoding="utf-8"))
CONS = json.load(open(os.path.join(HERE, "data", "consensus_spcx_20261030.json"), encoding="utf-8"))


def tk(name, code=None, case="基準"):
    v = TK["items"][name]["values"]
    return v if code is None else v[GEN[code]][case]


def hdr(ws, row, labels, c0=1):
    for j, h in enumerate(labels, c0):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR


def dn(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def fv(i): return F[i]["value"]
def fb(i): return F[i]["value"] / 1000.0                       # US$M → US$B
def fs(i):
    it = F[i]
    return f"{it['source']}（{it.get('docDate', '')}；事實總帳 {i}；{it['tag'].strip('[]')}）"


YK = ["25", "m25", "1H", "2H", "26", 27, 28, 29, 30, 31, 32]
COL = dict(zip(YK, "CDEFGHIJKLM"))
YH = {"25": "FY2025*", "m25": "9M25A", "1H": "9M26A", "2H": "Q4 26E", "26": "FY2026E", 27: "FY2027E", 28: "FY2028E", 29: "FY2029E",
      30: "FY2030E", 31: "FY2031E", 32: "FY2032E"}
FC = ["2H", 27, 28, 29, 30, 31, 32]
ANN = [27, 28, 29, 30, 31, 32]
HIST = ["25", "m25", "1H"]
LEN = {p: {"1H": 0.75, "m25": 0.75, "2H": 0.25}.get(p, 1) for p in YK}
PY = {"2H": 2026, 27: 2027, 28: 2028, 29: 2029, 30: 2030, 31: 2031, 32: 2032}
def prevp(p): return {"2H": "1H", 27: "2H", 28: 27, 29: 28, 30: 29, 31: 30, 32: 31}[p]
def prevfy(p): return {27: "26", 28: 27, 29: 28, 30: 29, 31: 30, 32: 31}[p]
MODEL_HDR = ["項目", "單位"] + [YH[k] for k in YK] + ["說明"]
NOTE_C = 14


def put(ws, r, label, unit, vals, fmt="#,##0.00", note=None, bold=False, fill=None):
    ws.cell(r, 1, label); ws.cell(r, 2, unit)
    if bold: ws.cell(r, 1).font = BOLD
    for p, v in vals.items():
        c = ws[f"{COL[p]}{r}"] if p in COL else ws[f"{p}{r}"]
        c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")) and isinstance(v, (int, float)):
            c.font = BLUE
    if note: ws.cell(r, NOTE_C, note)
    if fill:
        for cc in range(1, NOTE_C + 1): ws.cell(r, cc).fill = fill


def widths(ws, a=60, b=10, mid=11, note=90, ncols=11):
    ws.column_dimensions["A"].width = a; ws.column_dimensions["B"].width = b
    for i in range(ncols): ws.column_dimensions[chr(ord("C") + i)].width = mid
    ws.column_dimensions["N"].width = note


RD = wb.active; RD.title = "README"
IN = wb.create_sheet("In"); Q = wb.create_sheet("Q"); TH = wb.create_sheet("Thesis"); SG = wb.create_sheet("Seg")
PL = wb.create_sheet("PL"); CF = wb.create_sheet("CF"); V = wb.create_sheet("Val"); RC = wb.create_sheet("Recon")
SE = wb.create_sheet("Sens"); SO = wb.create_sheet("Sources")

# =====================================================================================================
# In：輸入（一列一鍵；E＝基準，D／F＝敏感度低／高，只用於 Sens 頁）
IN["A1"] = ("In — 輸入（藍字；基準欄 E 進主模型；D／F＝敏感度低／高，只用於 Sens 頁、不進目標價；情境 A／B／C 的驅動見本頁最後一節「情境驅動對照」）。"
            "金額 US$B；每股 US$；股數十億股；GW＝IT 口徑名目算力（不乘 PUE）。標記：Verified／Interested-party／Analogy／Assumed／Derived（Analogy、Assumed 一律給低／高區間）")
IN["A2"] = ("事實來源：spacex/data/spacex_facts_2026q3.json（10-Q 2026Q3 與 Q3／Q2 財報新聞稿兩讀核對；FY2025 為媒體轉述 S-1，標「二手轉述 S-1」）；"
            "Tokenomics v21.7 快照 spacex/data/tokenomics_snapshot_v21.7.json。黃底＝判斷題（Andy 可改；報告「待 Project 判斷」逐條列出預設）")
hdr(IN, 3, ["鍵", "項目", "單位", "敏感度低", "基準", "敏感度高", "標記", "來源", "說明"])
ROWS = []
def sec(t): ROWS.append(("sec", t))
def inp(key, item, unit, lo, base, hi, tag, src, note="", judg=False):
    ROWS.append((key, item, unit, lo, base, hi, tag, src, note, judg))

FQ = "mag/extract/docs/2026Q3/spacex_2026Q3_10-Q.pdf"
sec("A. 現價、股數與資產負債（評價日 2026-09-30；10-Q 2026Q3）")
inp("px", "現價（2026-10-29 收盤）", "US$", None, 143.89, None, "Verified",
    "MarketBeat SPCX 預測頁 https://www.marketbeat.com/stocks/NASDAQ/SPCX/forecast/（4:00 PM 收盤 143.89）；Equibles https://equibles.com/stocks/spcx/summary（10-29 收盤 143.89、跌 1.34；2026-10-30 讀取）",
    "兩來源一致（審查第 7 條）；Digrin 月線 10 月 143.89 亦同。A 段候選 Robinhood 140.40 為盤中價，不採")
inp("px_val", "評價日股價（2026-09-30 收盤；WACC 權益權重用）", "US$", None, 150.86, None, "Derived",
    "Digrin SPCX 月線 https://www.digrin.com/stocks/detail/SPCX/price（2026 年 9 月收盤 150.86；2026-10-30 讀取）",
    "單一來源（第二來源未取得：StockAnalysis 歷史頁逾時）；只影響 WACC 權重，負債占比 < 2%，對目標價影響極小")
inp("sh_out", "流通股數（10-Q 封面 2026-10-15，A＋B 股）", "十億股", None, fv("sh.out.AB.20261015") / 1000, None, "Verified", fs("sh.out.AB.20261015"),
    "A 8,246.5M＋B 5,334.5M（審查第 6 條；稀釋項列 Sens，不進基準）")
inp("sh_dil_9m", "9M26 稀釋加權股數", "十億股", None, fv("is.shDil.9m2026") / 1000, None, "Verified", fs("is.shDil.9m2026"), "IPO 在 2026-06，加權股數低於期末")
inp("dil_perf", "潛在稀釋：績效／市場條件獎酬（未計入稀釋股數）", "十億股", None, fv("sh.perfAwards.20260930") / 1000, None, "Verified", fs("sh.perfAwards.20260930"), "受獎人未見原文；Sens 用")
inp("dil_opt", "潛在稀釋：選擇權", "十億股", None, fv("sh.options.20260930") / 1000, None, "Verified", fs("sh.options.20260930"), "Sens 用")
inp("dil_rsu", "潛在稀釋：RSU", "十億股", None, fv("sh.rsu.20260930") / 1000, None, "Verified", fs("sh.rsu.20260930"), "Sens 用")
inp("dil_echo", "潛在稀釋：EchoStar 頻譜交割發行 A 股（約）", "十億股", None, fv("spectrum.echostar.shares") / 1000, None, "Verified", fs("spectrum.echostar.shares"), "固定每股 42.40；預計 2027-11-30 交割；Sens 用")
inp("dil_x", "評價股數加入上述潛在稀釋（1＝是；Sens 用）", "x", None, 0, 1, "Assumed", "審查第 6 條：稀釋項列 Sens，不進基準", "高＝四項全部加入（約 2.24B 股，＋16.5%）", judg=True)
inp("cash_0", "現金＋有價證券（2026-09-30）", "US$B", None, round(fb("bs.cash.20260930") + fb("bs.mktSec.20260930"), 3), None, "Verified",
    "10-Q 資產負債表：現金 51.294＋有價證券 38.561（事實總帳 bs.cash.20260930、bs.mktSec.20260930）", "6-30 為 100.009")
inp("cash_ye25", "現金＋有價證券（2025-12-31）", "US$B", None, round(fb("bs.cash.20251231") + fb("bs.mktSec.20251231"), 3), None, "Verified",
    "10-Q 資產負債表比較欄（事實總帳 bs.cash.20251231、bs.mktSec.20251231）")
inp("debt_0", "債務本金（2026-09-30；含失敗售後租回 12.897）", "US$B", None, fb("debt.totalDebt.principal.20260930"), None, "Verified", fs("debt.totalDebt.principal.20260930"),
    "SpaceX Notes 25.000＋其他融資 12.979（其中 Valor 失敗售後租回 12.897）＋X 票據 0.027")
inp("debt_ye25", "債務本金（2025-12-31）", "US$B", None, fb("debt.totalDebt.principal.20251231"), None, "Verified", fs("debt.totalDebt.principal.20251231"))
inp("slb_0", "其中：Valor 失敗售後租回債務（2026-09-30）", "US$B", None, fb("debt.failedSLB.total.20260930"), None, "Verified",
    "10-Q 附註 9、附註 17：直接寫明為 Valor Transactions（失敗售後租回；審查第 1 條），流動 2.181＋非流動 10.716", "關係人（董事 Antonio Gracias）；利息 2026Q3 0.351")
inp("fin_lease", "融資租賃負債（2026-09-30）", "US$B", None, fb("debt.finLease.principal.20260930"), None, "Verified", fs("debt.finLease.principal.20260930"), "計入 WACC 負債與淨負債")
inp("btc", "數位資產（比特幣 18,712 枚，2026-09-30 帳面）", "US$B", None, fb("bs.digital.20260930"), None, "Verified", fs("bs.digital.20260930"), "評價時加入（類現金持有）")
inp("legal", "訴訟準備（類債）", "US$B", None, fb("legal.accrual"), None, "Verified", fs("legal.accrual"))

sec("B. 折現率與目標價規則")
inp("rf", "無風險利率（10 年美債，2026-10-27）", "%", None, 0.0499, None, "Verified",
    "美國財政部每日殖利率曲線 10 Yr https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?type=daily_treasury_yield_curve&field_tdr_date_value_month=202610（Alphabet v0.3.2 同值，2026-10-29 讀取）")
inp("beta", "β（2026-06 上市，歷史不足；判斷）", "x", 1.1, 1.3, 1.6, "Assumed",
    "判斷：上市僅約 4 個月，無 3–5 年月資料 β；以高成長、高資本支出的大型科技股區間 1.1–1.6 取 1.3（Alphabet Blume 調整後 1.19 為下緣對照）", "Sens：1.1／1.6", judg=True)
inp("erp", "股權風險溢酬", "%", 0.04, 0.05, 0.06, "Assumed", "同 Alphabet v0.3.2（S9）")
inp("kd", "稅前債務成本（SpaceX Notes 有效利率）", "%", None, F["debt.spacexNotes.eir"]["value"] / 100, None, "Verified", fs("debt.spacexNotes.eir"), "五批加權票面 5.855%；有效 6.072%")
inp("tax_w", "WACC 用稅率（美國聯邦法定）", "%", None, 0.21, None, "Assumed", "淨營業虧損抵減期間實際稅盾低於法定；負債占比 < 2%，影響極小")
inp("wacc_adj", "WACC 調整（Sens 格：±1pt）", "%", -0.01, 0.0, 0.01, "Assumed", "Sens 用")
inp("g_tv", "永續成長率（FY2032 後）", "%", 0.02, 0.03, 0.04, "Assumed", "同 Alphabet v0.3.2（S9）", judg=True)
inp("ronic_sp", "永續期新投資報酬率 RONIC 相對 WACC 的差", "%", None, 0.0, 0.05, "Assumed", "同 Alphabet v0.3.2：RONIC＝WACC；終值＝NOPAT(FY2033) ×(1 − g ÷ RONIC) ÷ (WACC − g)", "Sens：WACC＋5pt", judg=True)
inp("mid", "期中慣例：年度現金流距年初", "年", None, 0.5, None, "Derived", "年度現金流平均發生在年中")
inp("t_val", "評價日（2026-09-30）的年份小數", "年", None, 2026.75, None, "Derived", "年度折現年數＝年度＋mid − t_val")
inp("t_q4", "Q4 26 現金流折現年數（期中）", "年", None, 0.125, None, "Derived", "2026-10-01～12-31 的中點")
inp("tp_wA", "情境權重 A", "%", None, 0.25, None, "Assumed", "工作單 S9（同 Alphabet：A 25%／B 50%／C 25%）", judg=True)
inp("tp_wB", "情境權重 B", "%", None, 0.50, None, "Assumed", "同上", judg=True)
inp("tp_wC", "情境權重 C", "%", None, 0.25, None, "Assumed", "同上", judg=True)
inp("w_dcf", "目標價中 DCF 的權重（其餘＝分部加總 SOTP）", "%", 0.3, 0.5, 0.7, "Assumed", "工作單 S9（DCF 50%＋分部加總 50%）", judg=True)
inp("t_tgt", "目標時點距評價日（2026-09-30 → 2027-10-30）", "年", None, round(395 / 365, 4), None, "Derived", "395 天 ÷ 365（工作單 S9：目標價時點 2027-10-30）")
inp("w_ntm27", "目標時點 NTM 中 FY2027 的權重", "%", None, round(62 / 365, 4), None, "Derived", "(2027-12-31 − 2027-10-30) ÷ 365＝62 天；FY2028 權重＝1 − 本值")
pt = CONS["priceTarget"]
inp("cons_tp", "市場共識目標價（平均）", "US$", pt["low"], pt["mean"], pt["high"], "Interested-party",
    f"{pt['source']} {pt['url']}（{pt['pageDate']}；{pt['nRatings']} 位分析師；data/consensus_spcx_20261030.json）", "賣方多有承銷關係，只作對照；低／高＝最低／最高目標價（只作對照，不進模型）")
inp("internal_ai_reading", "自用 AI 讀法（2＝主值：自用 AI 價值中性，計影子效益；1＝敏感度：只計成本）", "x", 1, 2, None, "Assumed",
    "原則第 7 節（Andy 2026-10-27 15:03 定案讀法 2）；工作單 S5：SI 自用 GW（＝在役 GW ×(1 − 對外出租占比)）評價日後新增者 × Tokenomics IF_HoldEcon（k＝1）",
    "營收列與分部對帳不動；Val 頁兩種讀法的目標價並列", judg=True)
inp("sbc_add", "DCF 加回股權報酬（1＝是；Sens 用）", "x", None, 0, 1, "Assumed", "同 Alphabet：股權報酬視為真實成本、不加回（含 Cursor 收購後服務期獎酬）", judg=True)
inp("tv_alt", "終值改用 FY2032 無槓桿自由現金流 ×(1＋g)（1＝是；Sens 用）", "x", None, 0, 1, "Assumed", "同 Alphabet", judg=True)

sec("C. 分部加總（SOTP）倍數（目標時點 NTM EBITDA × 同業 EV/EBITDA；Space 用 EV／營收；SI 對外出租用自身 DCF）")
for p_ in PCN["peers"]:
    if p_.get("evEbitdaFy2026") is None: continue
    inp(f"pm_conn_{p_['ticker'].lower()}", f"Connectivity 同業：{p_['name']}（{p_['ticker']}）EV/EBITDA FY2026", "x", None, p_["evEbitdaFy2026"], None, "Derived",
        f"{p_['source']} {p_['url']}（{p_['pageQuoteDate']}；data/peers_connectivity_20261030.json）", p_.get("note", ""))
for p_ in PSI["peers"]:
    x = p_.get("evEbitdaFy2027", p_.get("ntmEvEbitda"))
    lab = "EV/EBITDA FY2027" if "evEbitdaFy2027" in p_ else "NTM EV/EBITDA"
    inp(f"pm_si_{p_['ticker'].lower()}", f"SI 同業（AI 基礎設施／雲）：{p_['name']}（{p_['ticker']}）{lab}", "x", None, x, None, "Derived",
        f"{p_['source']}（{p_.get('pageQuoteDate', p_.get('asOf', ''))}；data/peers_si_20261030.json）", p_.get("note", ""))
inp("mult_cn_adj", "Connectivity 倍數＝同業中位數 × 本係數（Sens：同業區間下緣／上緣）", "x", 0.6606, 1.0, 2.3366, "Derived",
    "data/peers_connectivity_20261030.json：五家 FY2026 EV/EBITDA 區間 4.71（CMCSA）–16.66（SATS），中位數 7.13", "10-30 補 Sens：低／高＝區間下緣／上緣 ÷ 中位數")
inp("mult_cn_pm", "Connectivity 倍數 ±30%（Starlink 成長型對照倍數無依據時的替代）", "x", 0.7, 1.0, 1.3, "Assumed",
    "10-30 chat 端指示：Starlink 成長型對照倍數無可比上市公司依據（ASTS EBITDA 為負、NM）→ ±30%", "Sens 用")
inp("mult_si_adj", "自用 AI 增量營運成本加回的倍數＝SI 同業中位數 × 本係數（讀法 2；與 SI 其他估值拆開）", "x", 0.7, 1.0, 1.4, "Assumed",
    "審查第 6 條：同業樣本小（3 家、口徑不一），以中位數為基準、±30–40% 為區間；10-30 核對後只用於讀法 2 的自用營運成本加回", "", judg=True)
inp("evs_si", "SI 其他（Grok／Cursor／X 訂閱、資料授權、API 與廣告）EV／NTM 營收", "x", 5.0, 5.0, 8.0, "Assumed",
    "10-30 chat 端定（獨立核對建議）：殘差 EBITDA 率 16.9% 吸收 xAI 模型研發與自用機隊營運，不是業務經濟利潤率 → 改用 EV／NTM 營收，基準 5 倍、區間 5–8 倍",
    "對照：Cursor 收購對價 42.474（GAAP）約為其年化營收的 9 倍（Recon）；上市 AI 平台無可比公司（未找到）", judg=True)
inp("evs_space", "Space EV／NTM 營收", "x", 2.0, 3.0, 6.0, "Assumed",
    "工作單 S9：Space 無直接上市同業；Space 分部調整後 EBITDA 9M26 為負（Starship 開發），EBITDA 倍數不適用，改用 EV／營收（判斷）",
    "對照：Rocket Lab 等小型發射商 EV／營收約 10× 以上，但規模與獲利不可比（本次未重新查核）", judg=True)

sec("D. 命題模組：SI 名目算力（GW，IT 口徑）與資本支出（審查第 3 條）")
inp("gw_0", "SI 名目算力（2026-09-30）", "GW", None, fv("seg.si.kpi.gw.2026q3"), None, "Interested-party", fs("seg.si.kpi.gw.2026q3"),
    "10-Q 定義：安裝 GPU × all-in power，不含冷卻與設施負載（IT 口徑）；含對外出租、推定不含第三方雲（S6）")
inp("gw_q2", "SI 名目算力（2026-06-30）", "GW", None, fv("seg.si.kpi.gw.2026q2"), None, "Interested-party", fs("seg.si.kpi.gw.2026q2"))
inp("gw_ye25", "SI 名目算力（2025-12-31，內插）", "GW", 0.7, 0.8, 0.9, "Derived", "2025Q3 0.6 與 2026Q1 1.0 線性內插（年底值未揭露）", "汰換時點用")
inp("gw_ye24", "SI 名目算力（2024-12-31；汰換用）", "GW", 0.1, 0.15, 0.25, "Assumed", "判斷：Colossus 1 首批約 10 萬顆 H100（2024-09 上線，公司與媒體自述）≈ 0.15 GW IT；2025Q2 為 0.4", "未找到：2024 年底名目算力")
inp("ext_sh", "對外出租占比（SI S&I 營收與在役 GW 同比例；Customer B 衍生）", "%", 0.6, 0.72, 0.85, "Derived",
    "審查第 4 條：Customer B 2026Q3 占合併營收 34.5%（全屬 SI）→ 5.028 ÷ SI S&I 7.022＝71.6%（事實總帳 cust.B.rev.2026q3）；2026Q2 1.526 ÷ 2.194＝69.6%",
    "區間 0.6–0.85（審查第 4 條）；客戶 B 是否全為雲端服務屬推論；自用 GW＝在役 GW ×(1 − 本值)", judg=True)
inp("mix0_h", "評價日機隊：Hopper 錨（Colossus 1 H100／H200）", "%", None, 0.3, None, "Assumed", "判斷：Colossus 1 約 20 萬顆 Hopper（公司自述，二手）≈ 0.3–0.4 GW", "影響每 GW 持有成本的加權（k0 校準分母）")
inp("mix0_b2", "評價日機隊：GB200 錨", "%", None, 0.4, None, "Assumed", "判斷：Colossus 2 首批 GB200（2025H2–2026H1）")
inp("mix0_b3", "評價日機隊：GB300 錨", "%", None, 0.3, None, "Assumed", "判斷：Colossus 2 2026 年擴建（Q3 新聞稿：continued the buildout of COLOSSUS II）")
inp("am_27", "2027 新增中 GB300 錨占比（其餘 VR200 錨）", "%", None, 0.5, None, "Assumed", "同 Alphabet（2026 新增全為 GB300 錨；2028 起全為 VR200 錨）")
inp("selfb", "自建機房比例（不計機房資本支出的比例＝租用）", "%", 0.75, 0.9, 1.0, "Assumed", "判斷：Colossus 1／2 為自建（Memphis）；PP&E 資料中心基礎設施 6.333 偏低，可能部分租用", judg=True)
inp("si_cx_rr", "SI 投資季速（2026Q3 SI 資本支出＋以其他融資取得 PP&E 季均）", "US$B/季", None, round(fb("seg.si.capex.2026q3") + fb("cf.ppeFin.9m2026") / 3, 3), None, "Derived",
    "10-Q 附註 18：2026Q3 SI 資本支出 16.444＋以其他融資取得 PP&E 9M 3.921 ÷ 3（審查第 3 條：含以融資取得之 PP&E）", "以融資取得之 PP&E 未分分部，推定全屬 SI")
cxgw26 = tk("IF_CapexIT", "GB300") + 0.9 * tk("IF_CapexFacility", "VR200")
pace0 = round(4 * (fb("seg.si.capex.2026q3") + fb("cf.ppeFin.9m2026") / 3) / cxgw26, 4)
inp("pace", "SI 名目算力每年新增（FY2027–FY2028；三情境相同）", "GW/年", round(pace0 * 2 / 3, 4), "=4*si_cx_rr/(cit_b3+selfb*cfac)", round(pace0 * 1.5, 4), "Derived",
    "審查第 3 條：SI 投資季速 × 4 ÷ Tokenomics 每 GW 資本支出（GB300 錨 IT＋自建 × 機房）", f"基準約 {pace0:.3f}；Q4 26 新增＝本值 × 0.25（年底約 2.1 GW，公司 Q2 法說稱逾 2 GW，Recon 對照）；低／高＝× 2/3／× 1.5")
inp("conv_mw", "FY2029 起新增速度收斂年數（到 FY2032 只剩 g × 期初 GW）", "年", None, 4, None, "Assumed", "同 Alphabet 修正一 F5：終值不得落在建置高峰", judg=True)
inp("fin_sh", "SI 投資中以其他融資取得（非現金）比例", "%", None, round(fb("cf.ppeFin.9m2026") / (fb("seg.si.capex.9m2026") + fb("cf.ppeFin.9m2026")), 4), None, "Derived",
    "9M26：3.921 ÷ (SI 資本支出 39.995＋3.921)", "預測期 SI 投資按本比例拆現金資本支出與融資取得；收支關卡兩者都扣（審查第 2 條）")
for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3"), ("VR200", "vr")):
    inp(f"he_{nm}", f"IF_HoldEcon｜{code}（每 MW 年持有成本，經濟口徑）", "US$m/MW-年", tk("IF_HoldEcon")[GEN[code]]["低成本"], tk("IF_HoldEcon", code), tk("IF_HoldEcon")[GEN[code]]["高成本"], "Derived",
        f"data/tokenomics_snapshot_{TKV}.json（{TK['source']['file']}，{TK['items']['IF_HoldEcon']['ref']}）", "Tokenomics WACC 10% 下的打平持有成本；＝US$B/GW-年")
for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3"), ("VR200", "vr")):
    inp(f"opx_{nm}", f"IF_OpexGW｜{code}（每 MW 年營運費用，不含折舊）", "US$m/MW-年", None, tk("IF_OpexGW", code), None, "Derived", f"data/tokenomics_snapshot_{TKV}.json")
for code, nm in (("GB300", "b3"), ("VR200", "vr")):
    inp(f"cit_{nm}", f"IF_CapexIT｜{code}（每 MW IT 資本支出）", "US$m/MW", None, tk("IF_CapexIT", code), None, "Derived", f"data/tokenomics_snapshot_{TKV}.json")
inp("cfac", "IF_CapexFacility（每 MW 機房資本支出）", "US$m/MW", None, tk("IF_CapexFacility", "VR200"), None, "Derived", f"data/tokenomics_snapshot_{TKV}.json")
for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3")):
    inp(f"dpi_{nm}", f"IF_DeprIT｜{code}（每 MW IT 年折舊）", "US$m/MW-年", None, tk("IF_DeprIT", code), None, "Derived", f"data/tokenomics_snapshot_{TKV}.json", "期初機隊汰換時遞減基礎折舊")
inp("life_it", "AI IT 折舊年限", "年", 5, 6, 7, "Derived", "Tokenomics IF_DeprLifeIT 6 年")
inp("life_fac", "機房折舊年限", "年", 15, 20, 25, "Derived", "Tokenomics IF_CapexFacility 12.67 ÷ IF_DeprFac 0.634＝20 年")
inp("phi_0", "資本支出當期折舊比例（在役時點與半年慣例）", "x", 0.0, 0.25, 0.5, "Assumed", "同 Alphabet")
inp("cip_0", "在建工程（2026-09-30）", "US$B", None, fb("ppe.cip.20260930"), None, "Verified", fs("ppe.cip.20260930"), "主要為伺服器與資料中心")
inp("cip_si", "在建工程中屬 SI 的比例", "%", 0.0, 0.85, 1.0, "Assumed", "判斷：9M26 SI 占資本支出 83%（39.995 ÷ 48.069）", "SI 部分 Q4 26 起依 IT／機房年限折舊（不另增 GW；同 Alphabet 修正三口徑，判斷題）", judg=True)
inp("cip_it", "在建工程（SI）中 IT 占比", "%", None, 0.8, None, "Assumed", "判斷：9M26 伺服器與網通 +20.5、資料中心基礎設施 +3.4（PP&E 附註）")
inp("da_si0", "SI 評價日在役資產年折舊（2026Q3 × 4）", "US$B/年", None, round(fb("seg.si.da.2026q3") * 4, 3), None, "Derived", fs("seg.si.da.2026q3"), "期初機隊汰換時遞減（同 Alphabet 修正一）")
inp("it_sh_da", "SI 基礎折舊中 IT 占比（汰換遞減上限）", "%", None, 0.8, None, "Assumed", "判斷：伺服器 43.2 對資料中心 6.3（PP&E 附註）")

sec("E. 命題模組：SI 每 GW 年營收（審查第 4 條；情境軸）")
inp("k30_A", "情境 A：FY2030 每 GW 年營收 ÷ Tokenomics 持有成本（k）", "x", None, 1.6, None, "Assumed",
    "判斷：維持 2026Q3 校準值附近（供給吃緊、對外合約三年固定月費；Q2 新簽雲端合約 14.1B）", "FY2026Q4＝校準 k0，線性到 FY2030，之後持平", judg=True)
inp("k30_B", "情境 B：FY2030 每 GW 年營收 ÷ Tokenomics 持有成本（k）", "x", 1.0, 1.25, 1.5, "Assumed",
    "判斷：MAG 長約與現貨證據混合 k 約 1.2（Alphabet v0.3.2 Thesis k_blend）；合約到期後重新議價向持有成本收斂", "", judg=True)
inp("k30_C", "情境 C：FY2030 每 GW 年營收 ÷ Tokenomics 持有成本（k）", "x", None, 1.0, None, "Assumed",
    "判斷：算力供給追上後租金收斂到打平持有成本（報酬＝Tokenomics WACC 10%）", "", judg=True)
inp("conv_k", "k 收斂年數（Q4 26 校準值 → 情境值）", "年", 2, 4, 6, "Assumed", "同成長率收斂（FY2030 到位）",
    "Sens：2／6 年；另列「k 不收斂」一案（conv_k 極大，k 維持 k0）。注意：k0 隨 Tokenomics 持有成本重新校準、FY2030 k 為固定值，故持有成本同時推動營收")

sec("F. SI 其他收入（Grok／Cursor／X 訂閱、資料授權、API；X 廣告）")
inp("cur_pf", "Cursor 備考補足（2026Q3 備考營收 − 實際營收）", "US$B/季", None, round((fv("cursor.proforma.rev.2026q3") - fv("is.rev.2026q3")) / 1000, 3), None, "Interested-party",
    fs("cursor.proforma.rev.2026q3"), "審查第 8 條：Cursor 2026-08-14 併入，Q3 只含約 1.5 月（608）；預測期以備考全季為起點（未經查核）")
inp("g27_nr", "SI 其他收入 FY2027 年增（以 Q4 26 備考年化為基期）", "%", 0.25, 0.5, 0.8, "Assumed",
    "判斷：Grok 訂閱與 API、Cursor（AI 程式助理）高速成長；公司未揭露拆分（A 段第 5 節）", "", judg=True)
inp("gl_nr", "SI 其他收入長期成長（FY2030 到位）", "%", 0.08, 0.15, 0.25, "Assumed", "判斷", "", judg=True)
inp("gl_ads", "X 廣告長期成長", "%", -0.03, 0.03, 0.08, "Assumed", "判斷：9M26 年減 5.7%、2026Q3 年增 18.5%（Q 頁）")
inp("conv_n", "成長率收斂年數（起始 → 長期值）", "年", None, 4, None, "Assumed", "同 Alphabet（FY2030 到達）")
inp("m_so_adj", "SI 其他 EBITDA 率調整（加在 2026Q3 校準值上）", "%", -0.1, 0.0, 0.1, "Assumed", "Thesis 頁 m_so0：2026Q3 SI 調整後 EBITDA 扣對外出租 EBITDA 後的殘差率", "Sens 用")
inp("sbc_cur_left", "Cursor 收購後服務期股權報酬：尚未認列（2026-09-30）", "US$B", None, round((fv("cursor.postSvcAwards") - fv("seg.si.sbcAcq.2026q3")) / 1000, 3), None, "Derived",
    "10-Q 附註：收購後服務期獎酬 23.376，約 2.8 年認列；2026Q3 已認列收購相關 SBC 1.217（事實總帳 cursor.postSvcAwards、seg.si.sbcAcq.2026q3）", "直線攤到 2029-06（Q4 26 起每年約 8.4）；非現金、DCF 不加回")
inp("sbc_cur_yrs", "Cursor 獎酬剩餘認列年數（2026-10-01 起）", "年", None, 2.65, None, "Derived", "2026-08-14 起 2.8 年 − 已過 0.13 年")

sec("G. Space（發射次數 × 每次收入＋發射與開發）")
inp("nl_q4", "Q4 26 客戶發射次數", "次", None, 12, None, "Assumed", "2026Q3 實際 12、Q2 10、Q1 7（Q 頁）")
inp("nl_27", "FY2027 客戶發射次數", "次", 40, 48, 60, "Assumed", "判斷：Q3 年化 48；2025 年 9M 36 次", judg=True)
inp("g_nl", "客戶發射次數年增（FY2028 起）", "%", 0.0, 0.05, 0.1, "Assumed", "判斷")
inp("g_rpl", "每次客戶發射收入年增", "%", 0.0, 0.03, 0.06, "Assumed", "判斷：起點＝9M26 發射服務 2.108 ÷ 29 次（Q 頁）")
inp("gl_ld", "發射與開發（政府開發合約等）長期成長", "%", 0.0, 0.08, 0.15, "Assumed", "判斷：起點＝9M 年增（Q 頁）；HLS 等里程碑收入不規則")
inp("m_spl", "Space 分部調整後 EBITDA 率 FY2030 到達", "%", 0.05, 0.2, 0.3, "Assumed", "判斷：9M26 −5.9%（Starship 開發）；FY2025 二手轉述 16.0%", judg=True)
inp("cx_spl", "Space 資本支出占營收 FY2030 到達", "%", 0.2, 0.3, 0.5, "Assumed", "判斷：9M26 114%（Starship 量產與發射台）", judg=True)
inp("life_sp", "Space 新增資本支出折舊年限", "年", None, 10, None, "Assumed", "判斷：發射台、工廠與可重複使用載具")
inp("rem_sp", "Space 評價日在役資產剩餘折舊年數（基礎折舊直線遞減到 0）", "年", None, 5, None, "Assumed", "判斷")

sec("H. Connectivity（消費訂戶 × ARPU＋企業與政府）")
inp("nadd_q4", "Q4 26 消費訂戶淨增", "百萬", None, round(fv("seg.conn.kpi.subs.2026q3") - fv("seg.conn.kpi.subs.2026q2"), 2), None, "Interested-party",
    "2026Q3 淨增 13.6 − 12.0＝1.6（新聞稿 KPI）")
inp("nadd_A", "情境 A：FY2027 消費訂戶淨增", "百萬", None, 7.0, None, "Assumed", "判斷：Q3 年化 6.4；直連手機（EchoStar 頻譜）與新興市場", judg=True)
inp("nadd_B", "情境 B：FY2027 消費訂戶淨增", "百萬", 4.0, 5.5, 7.0, "Assumed", "判斷：Q3 年化 6.4，淨增逐步放緩", judg=True)
inp("nadd_C", "情境 C：FY2027 消費訂戶淨增", "百萬", None, 4.0, None, "Assumed", "判斷：低價方案帶來的增長放緩、競爭（Kuiper）", judg=True)
inp("nadd_dec", "FY2028 起淨增每年遞減", "%", None, 0.1, None, "Assumed", "判斷")
inp("arpu0", "消費 ARPU（2026Q3）", "US$/月", None, fv("seg.conn.kpi.arpu.2026q3"), None, "Interested-party", fs("seg.conn.kpi.arpu.2026q3"), "2025Q2 85 → 2026Q1 66 → Q3 69")
inp("g_arpu", "ARPU 年變動", "%", -0.05, -0.02, 0.0, "Assumed", "判斷：新興市場與低價方案占比上升（2025 年 ARPU 由 85 降到 66）")
inp("g0_eg_adj", "企業與政府：起始年增調整（加在 9M 年增上）", "%", None, 0.0, None, "Assumed", "起點＝9M26 年增（Q 頁）")
inp("gl_eg", "企業與政府長期成長", "%", 0.05, 0.1, 0.15, "Assumed", "判斷：航空、海事、Starshield 政府合約")
inp("m_cn", "Connectivity 分部調整後 EBITDA 率（預測期）", "%", 0.58, "=Q!B{mcn}", 0.7, "Derived", "9M26 Connectivity 調整後 EBITDA 8.183 ÷ 營收 12.796（Q 頁）")
inp("cx_cn", "Connectivity 資本支出占營收", "%", 0.25, "=Q!B{cxcn}", 0.4, "Derived", "9M26 4.242 ÷ 12.796（Q 頁）")
inp("life_cn", "Connectivity 新增資本支出折舊年限（衛星）", "年", None, 6, None, "Assumed", "判斷：Starlink V2 mini 設計壽命約 5 年")
inp("rem_cn", "Connectivity 評價日在役資產剩餘折舊年數", "年", None, 3, None, "Assumed", "判斷")
inp("spec_2h", "EchoStar 頻譜信用協議付款 Q4 26", "US$B", None, round((fv("spectrum.echostar.pay2026") - fv("spectrum.echostar.paid")) / 1000, 3), None, "Verified",
    "10-Q MD&A Spectrum Transaction：2026 年 1.241、已付 0.856（事實總帳 spectrum.echostar.pay2026、.paid）", "投資活動；頻譜按成本計（價值中性）")
inp("spec_27", "EchoStar 頻譜：FY2027 現金（清償指定債務最多 8.5＋信用協議 0.828）", "US$B", None, round(fv("spectrum.echostar.cash") + fv("spectrum.echostar.pay2027") / 1000, 3), None, "Verified",
    "10-Q MD&A：清償 EchoStar 指定債務最多 8.5（不足以現金補）、2027 年付款 0.828；預計 2027-11-30 交割", "股權部分（約 261.8M 股）列 Sens（dil_echo）；800 MHz 交易金額未揭露（未找到）")

sec("I. 損益與現金流")
inp("tax", "有效稅率（明確期 Q4 26–FY2032）", "%", 0.05, 0.12, 0.21, "Assumed", "判斷：累積虧損與淨營業虧損抵減；2026Q3 認列稅務利益（含評價備抵迴轉 0.449）", "", judg=True)
inp("tax_tv", "終值期稅率（FY2033 起）", "%", 0.12, 0.21, None, "Assumed", "10-30 chat 端定：淨營業虧損用完後回到美國聯邦法定 21%", "Sens 低＝12% 永續（舊口徑）", judg=True)
inp("r_cash", "現金收益率", "%", None, round(fb("is.intInc.2026q3") * 4 / ((fb("bs.cash.20260630") + fb("bs.mktSec.20260630") + fb("bs.cash.20260930") + fb("bs.mktSec.20260930")) / 2), 4), None, "Derived",
    "2026Q3 利息收入 0.873 × 4 ÷ 平均現金與有價證券（6-30 100.0、9-30 89.9）")
inp("r_debt", "債務利率（含失敗售後租回）", "%", None, round(-fb("is.intExp.2026q3") * 4 / fb("debt.totalDebt.principal.20260930"), 4), None, "Derived", "2026Q3 利息費用 0.675 × 4 ÷ 債務本金 38.0")
inp("dil", "股數年稀釋（股權報酬）", "%", 0.0, 0.01, 0.02, "Assumed", "同 Alphabet 模板每年 1%")
inp("wc_k", "營運資金變動 ÷ 營收增量（流入）", "%", 0.0, round(sum(fv(f"cf.wc.{k}.9m2026") for k in ("ar", "inv", "prepaid", "ap", "dr", "oth")) / (fv("seg.total.rev.9m2026") - fv("seg.total.rev.9m2025")), 4), 0.1, "Derived",
    "9M26 營運資金六列合計 0.375 ÷ 營收增量 13.369（遞延收入 +4.9 被應收、存貨抵銷）")
inp("min_cash", "最低現金", "US$B", None, 30, None, "Assumed", "判斷：約一季資本支出")
for p, v in (("2H", fb("debt.maturity.2026Q4")), (27, fb("debt.maturity.2027")), (28, fb("debt.maturity.2028")), (29, fb("debt.maturity.2029")), (30, fb("debt.maturity.2030")), (31, 5.0), (32, 5.0)):
    inp(f"mat_{p}", f"債務到期（推定主要為 Valor 失敗售後租回本金）{YH[p]}", "US$B", None, v, None, "Verified" if p not in (31, 32) else "Assumed",
        "10-Q 附註 9 到期表；SpaceX Notes 2031-07-15 前無重大本金（MD&A）→ 2026Q4–2030 推定主要為失敗售後租回與設備融資攤還" if p == "2H" else ("FY2031 起 5.0／年（其後 25.2 未分年）" if p == 31 else ""))
for p, v in (("2H", fb("debt.finLease.principal.20260930") / 8), (27, 0.25), (28, 0.2), (29, 0.15), (30, 0.1), (31, 0.1), (32, 0.1)):
    inp(f"lease_{p}", f"融資租賃本金 {YH[p]}", "US$B", None, round(v, 3), None, "Assumed", "判斷：融資租賃負債 0.983；9M26 本金 0.270" if p == "2H" else "")

sec("I2. 2026 年 9M 實績（CF 頁；10-Q 現金流量表）")
inp("cfo_9m", "9M26 營運現金流", "US$B", None, fb("cf.cfo.9m2026"), None, "Verified", fs("cf.cfo.9m2026"))
inp("capex_9m", "9M26 購置不動產與設備（＝分部資本支出合計）", "US$B", None, -fb("cf.capex.9m2026"), None, "Verified", fs("cf.capex.9m2026"))
inp("ppefin_9m", "9M26 以其他融資取得 PP&E（非現金，未入資本支出）", "US$B", None, fb("cf.ppeFin.9m2026"), None, "Verified", fs("cf.ppeFin.9m2026"))
inp("lease_1H", "9M26 融資租賃本金", "US$B", None, -fb("cf.flPrin.9m2026"), None, "Verified", fs("cf.flPrin.9m2026"))
inp("slb_1H", "9M26 失敗售後租回本金償還", "US$B", None, 0.0, None, "Derived", "未單獨揭露（含在償還債務 39.845；主要為過橋貸款與 xAI 舊債清償）→ 0", "資料缺口")
inp("debt_1H", "9M26 債務淨流入（發行 − 償還 − 發行成本 − 清償溢價）", "US$B", None,
    round(fb("cf.debtIss.9m2026") + fb("cf.debtRep.9m2026") + fb("cf.debtCost.9m2026") + fb("cf.extPrem.9m2026"), 3), None, "Verified",
    "10-Q 現金流量表：發行 51.834 − 償還 39.845 − 成本 − 清償溢價 1.153")
inp("eq_1H", "9M26 股權淨流入（IPO 淨額＋IPO 前發行＋員工認股 − 買回 − 代扣稅）", "US$B", None,
    round(fb("cf.ipo.9m2026") + fb("cf.stockIss.9m2026") + fb("cf.espp.9m2026") + fb("cf.buyback.9m2026") + fb("cf.taxNet.9m2026"), 3), None, "Verified",
    "10-Q 現金流量表：IPO 淨額 85.675、IPO 前發行 8.319、買回 −4.426 等")
inp("sh_dil_1H", "9M26 稀釋加權股數（EPS 用）", "十億股", None, fv("is.shDil.9m2026") / 1000, None, "Verified", fs("is.shDil.9m2026"))
inp("deemed_1H", "9M26 視同股利（特別股轉換；歸屬普通股淨利扣除）", "US$B", None, fb("is.deemedDiv.9m2026"), None, "Verified", fs("is.deemedDiv.9m2026"))

sec("J0. 對照（Recon）")
inp("gw_ye26_co", "公司稱 2026 年底名目算力（Q2 法說，二手轉述）", "GW", None, 2.0, None, "Interested-party", "審查第 3 條：公司 Q2 法說稱年底逾 2 GW（二手轉述、Interested-party）", "只作 Recon 對照，不進模型")
inp("backlog", "Backlog（2026-09-30，未分分部）", "US$B", None, fb("backlog.total.20260930"), None, "Verified", fs("backlog.total.20260930"), "67% 一年內、27% 一至三年")
inp("cloud_q2", "2026Q2 新簽雲端服務合約（合約金額）", "US$B", None, fb("si.cloud.contracted.2026q2"), None, "Interested-party", fs("si.cloud.contracted.2026q2"))
inp("commit_27", "不可取消承諾 FY2027（SI 基礎設施、第三方雲容量、頻譜等）", "US$B", None, fb("commit.uncond.2027"), None, "Verified", fs("commit.uncond.2027"), "未拆分算力採購與頻譜")
SWS = CONS["annual"]["stale"]
for y in ("2026", "2027", "2028"):
    for k_, lab in (("revenue", "營收"), ("earnings", "淨利"), ("fcf", "自由現金流")):
        inp(f"c_{k_}_{y[2:]}", f"過期共識 {lab} FY{y}（Simply Wall St，2026-06-22）", "US$B", None, SWS[f"FY{y}"][k_] / 1000, None, "Interested-party",
            SWS["source"] + "（網址見 data/consensus_spcx_20261030.json）" if (y == "2026" and k_ == "revenue") else "同上", "Q3 後年度共識未找到（A 段已搜尋 MarketBeat、StockAnalysis、S&P MI、Chartmill）；只作方向對照" if (y == "2026" and k_ == "revenue") else "")

# 寫入 In
IN_ROW = {}
r = 4
for row in ROWS:
    if row[0] == "sec":
        IN.cell(r, 1, row[1]).font = BOLD; r += 1; continue
    key = row[0]
    IN_ROW[key] = r; r += 1
r = 4
for row in ROWS:
    if row[0] == "sec":
        r += 1; continue
    key, item, unit, lo, base, hi, tag, src, note, judg = row
    if key in OVR: base = OVR[key]
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            c = IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
            c.number_format = "0.0%" if unit == "%" else ("0.0000" if unit == "x" else "#,##0.000")
    IN.cell(r, 7, tag); IN.cell(r, 8, src); IN.cell(r, 9, note)
    if judg:
        for cc in range(1, 10): IN.cell(r, cc).fill = KEY
    dn(key, f"In!$E${r}")
    r += 1
SCN_ROWS = [("scn_k", "SI 每 GW 年營收 ÷ Tokenomics 持有成本（k）FY2030 到達", "x", ("k30_A", "k30_B", "k30_C"), "0.00", "Q4 26 起由 2026Q3 校準值 k0 線性收斂"),
            ("scn_n", "Starlink 消費訂戶 FY2027 淨增", "百萬", ("nadd_A", "nadd_B", "nadd_C"), "0.0", "FY2028 起每年遞減 nadd_dec"),
            ("scn_w", "情境權重", "%", ("tp_wA", "tp_wB", "tp_wC"), "0%", "[判斷] 黃底，Andy 可改")]
r += 1
IN.cell(r, 1, "J. 情境驅動對照（三情境並排，進主模型的加權目標價；本節只引用上方各列，修改請改上方藍字。上方「敏感度低／高」欄只用於 Sens 頁，不進目標價）").font = BOLD; r += 1
hdr(IN, r, ["鍵", "項目", "單位", "情境 A（樂觀）", "情境 B（基準）", "情境 C（保守）", "標記", "說明"]); r += 1
for key, item, unit, ks, fmt, note in SCN_ROWS:
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, k_ in zip((4, 5, 6), ks):
        IN.cell(r, j, f"=E{IN_ROW[k_]}" if k_ else 0).number_format = fmt
    IN.cell(r, 7, "Derived"); IN.cell(r, 8, "引用 In 頁 " + "／".join(k_ for k_ in ks if k_) + "；" + note)
    r += 1
for col, w in zip("ABCDEFGHI", [14, 64, 12, 10, 11, 10, 16, 90, 80]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# =====================================================================================================
# Q：實績（2025Q2–2026Q3、9M25、9M26、FY2025*）
Q["A1"] = ("Q — 實績（US$B；EPS 美元；股數十億股；GW、訂戶、次數為原單位）。資料取自 data/spacex_facts_2026q3.json（10-Q 2026Q3、Q3 新聞稿、Q2 新聞稿兩讀核對）；"
           "FY2025* 為媒體轉述 S-1（二手，待 424B4 原文；Space 資本支出未取得）")
Q["A2"] = "核對列（實績必須為 0）：子列合計＝分部、分部合計＝合併（營收、營業利益、分部調整後 EBITDA、資本支出）、分部調整後 EBITDA 恆等式、損益表恆等式"
QCOLS = ["2025q2", "2025q3", "2026q1", "2026q2", "2026q3", "9m2025", "9m2026", "fy2025"]
QH = ["2025Q2", "2025Q3", "2026Q1", "2026Q2", "2026Q3", "9M25", "9M26", "FY2025*"]
QC = dict(zip(QCOLS, "BCDEFGHI"))
hdr(Q, 3, ["項目"] + QH + ["標記", "來源"])
QR = {}
SRCQ = "10-Q 2026Q3 附註 3、18（四期）；Q3 新聞稿（2026Q2 欄）；Q2 新聞稿（2026Q1、2025Q2）；FY2025＝媒體轉述 S-1（事實總帳 seg.*／is.*）"
PER5 = ["2025q3", "2026q2", "2026q3", "9m2025", "9m2026"]


def qrow(r, key, label, fid=None, scale=1000, neg=False, formula=None, fmt="#,##0.000", tag="Verified／Interested-party（KPI）", src=SRCQ, cols=None):
    Q.cell(r, 1, label); QR[key] = r
    for c in (cols or QCOLS):
        col = QC[c]
        if formula:
            Q[f"{col}{r}"] = "=" + formula(col)
        else:
            iid = f"{fid}.{c}"
            if iid not in F or F[iid]["value"] is None: continue
            v = F[iid]["value"]
            Q[f"{col}{r}"] = (-v if neg else v) / scale
            Q[f"{col}{r}"].font = BLUE
        Q[f"{col}{r}"].number_format = fmt
    Q.cell(r, 10, tag); Q.cell(r, 11, src)


r = 4
SEGS = (("space", "Space"), ("conn", "Connectivity"), ("si", "SI（原 AI）"))
Q.cell(r, 1, "營收（官方分部與子列）").font = BOLD; r += 1
NOFY = [c for c in QCOLS if c != "fy2025"]
qrow(r, "launchSvc", "Launch Services", "seg.space.launchSvc"); r += 1
qrow(r, "launchDev", "Launch & Development", "seg.space.launchDev"); r += 1
qrow(r, "space", "Space（分部）", "seg.space.rev"); r += 1
qrow(r, "chk_sp", "檢查：Space 子列合計 − Space", formula=lambda c: f"{c}{QR['launchSvc']}+{c}{QR['launchDev']}-{c}{QR['space']}", fmt="0.000", tag="—", src="—", cols=NOFY); r += 1
qrow(r, "consumer", "Consumer", "seg.conn.consumer"); r += 1
qrow(r, "entGov", "Enterprise & Government", "seg.conn.entGov"); r += 1
qrow(r, "conn", "Connectivity（分部）", "seg.conn.rev"); r += 1
qrow(r, "chk_cn", "檢查：Connectivity 子列合計 − Connectivity", formula=lambda c: f"{c}{QR['consumer']}+{c}{QR['entGov']}-{c}{QR['conn']}", fmt="0.000", tag="—", src="—", cols=NOFY); r += 1
qrow(r, "ads", "Advertising", "seg.si.ads"); r += 1
qrow(r, "ssi", "SI Solutions & Infrastructure", "seg.si.siSI"); r += 1
qrow(r, "si", "SI（分部）", "seg.si.rev"); r += 1
qrow(r, "chk_si", "檢查：SI 子列合計 − SI", formula=lambda c: f"{c}{QR['ads']}+{c}{QR['ssi']}-{c}{QR['si']}", fmt="0.000", tag="—", src="—", cols=NOFY); r += 1
qrow(r, "rev", "合併營收", "seg.total.rev"); r += 1
qrow(r, "chk_rev", "檢查：分部合計 − 合併營收", formula=lambda c: f"{c}{QR['space']}+{c}{QR['conn']}+{c}{QR['si']}-{c}{QR['rev']}", fmt="0.000", tag="—", src="—"); r += 1
r += 1
for key, nm in (("ebit", "分部營業利益（損失）"), ("adjEbitda", "分部調整後 EBITDA（公司定義，不含 SBC）"), ("capex", "分部資本支出（購置不動產與設備）")):
    Q.cell(r, 1, nm).font = BOLD; r += 1
    for s, lab in SEGS:
        qrow(r, f"{key}_{s}", f"{lab}：{nm[:nm.index('（') if '（' in nm else len(nm)]}", f"seg.{s}.{key}"); r += 1
    qrow(r, f"{key}_tot", f"合計：{nm[:nm.index('（') if '（' in nm else len(nm)]}", f"seg.total.{key}"); r += 1
    cols = [c for c in QCOLS if not (c == "fy2025" and key == "capex")]
    qrow(r, f"chk_{key}", f"檢查：分部合計 − 合併（{nm[:nm.index('（') if '（' in nm else len(nm)]}）",
         formula=lambda c, key=key: f"{c}{QR[key + '_space']}+{c}{QR[key + '_conn']}+{c}{QR[key + '_si']}-{c}{QR[key + '_tot']}", fmt="0.000", tag="—", src="—", cols=cols); r += 1
    r += 1
Q.cell(r, 1, "分部調整後 EBITDA 調節項（折舊攤銷、股權報酬、收購相關、重組、減損）").font = BOLD; r += 1
for s, lab in SEGS:
    qrow(r, f"da_{s}", f"{lab}：折舊攤銷", f"seg.{s}.da", cols=NOFY); r += 1
    qrow(r, f"sbc_{s}", f"{lab}：股權報酬（一般）", f"seg.{s}.sbc", cols=NOFY); r += 1
qrow(r, "sbcAcq_si", "SI：收購相關股權報酬（Cursor）", "seg.si.sbcAcq", cols=NOFY); r += 1
qrow(r, "acq_si", "SI：收購成本", "seg.si.acqCost", cols=NOFY); r += 1
qrow(r, "restr_si", "SI：重組費用（X）", "seg.si.restr", cols=NOFY); r += 1
qrow(r, "imp_sp", "Space：減損", "seg.space.impair", cols=NOFY); r += 1
for s, lab in SEGS:
    extra = {"space": f"+N({{c}}{QR['imp_sp']})", "conn": "", "si": f"+N({{c}}{QR['sbcAcq_si']})+N({{c}}{QR['acq_si']})+N({{c}}{QR['restr_si']})"}[s]
    qrow(r, f"chk_e_{s}", f"檢查：{lab} 調整後 EBITDA − (營業利益＋折舊攤銷＋SBC＋其他調節)",
         formula=lambda c, s=s, extra=extra: f"{c}{QR['adjEbitda_' + s]}-({c}{QR['ebit_' + s]}+{c}{QR['da_' + s]}+{c}{QR['sbc_' + s]}" + extra.replace("{c}", c) + ")",
         fmt="0.000", tag="—", src="五期（10-Q／Q3 新聞稿）", cols=PER5); r += 1
r += 1
Q.cell(r, 1, "營運指標").font = BOLD; r += 1
qrow(r, "nlc", "客戶發射（次）", "seg.space.kpi.launchCust", scale=1, fmt="0", cols=NOFY); r += 1
qrow(r, "nlt", "發射合計（次）", "seg.space.kpi.launchTot", scale=1, fmt="0", cols=NOFY); r += 1
qrow(r, "subs", "Starlink 消費訂戶（百萬，期末）", "seg.conn.kpi.subs", scale=1, fmt="0.0", cols=NOFY); r += 1
qrow(r, "arpu", "消費 ARPU（美元／月）", "seg.conn.kpi.arpu", scale=1, fmt="0", cols=NOFY); r += 1
qrow(r, "gw", "SI 名目算力（GW，期末；IT 口徑）", "seg.si.kpi.gw", scale=1, fmt="0.0", cols=NOFY); r += 1
r += 1
Q.cell(r, 1, "合併損益（10-Q 四期；2026Q2 只有新聞稿 Highlights）").font = BOLD; r += 1
IS4 = ["2025q3", "2026q3", "9m2025", "9m2026"]
qrow(r, "costs", "成本費用合計", "is.costs", cols=IS4); r += 1
qrow(r, "is_ebit", "營業利益", "is.ebit", cols=IS4 + ["2026q2"]); r += 1
qrow(r, "chk_cost", "檢查：營收 − 成本費用 − 營業利益", formula=lambda c: f"{c}{QR['rev']}-{c}{QR['costs']}-{c}{QR['is_ebit']}", fmt="0.000", tag="—", src="—", cols=IS4); r += 1
qrow(r, "intExp", "利息費用（含 Valor 關係人）", "is.intExp", cols=IS4); r += 1
qrow(r, "intInc", "利息收入", "is.intInc", cols=IS4); r += 1
qrow(r, "oie", "其他收益（費用）淨額", "is.oie", cols=IS4); r += 1
qrow(r, "pt", "稅前利益", "is.pretax", cols=IS4); r += 1
qrow(r, "chk_pt", "檢查：營業利益＋利息＋其他 − 稅前", formula=lambda c: f"{c}{QR['is_ebit']}+{c}{QR['intExp']}+{c}{QR['intInc']}+{c}{QR['oie']}-{c}{QR['pt']}", fmt="0.000", tag="—", src="—", cols=IS4); r += 1
qrow(r, "taxa", "所得稅費用（負值＝利益）", "is.tax", cols=IS4); r += 1
qrow(r, "ni", "淨利（損）", "is.ni", cols=IS4 + ["2026q2"]); r += 1
qrow(r, "chk_ni", "檢查：稅前 − 所得稅 − 淨利", formula=lambda c: f"{c}{QR['pt']}-{c}{QR['taxa']}-{c}{QR['ni']}", fmt="0.000", tag="—", src="—", cols=IS4); r += 1
qrow(r, "nic", "歸屬普通股淨利（扣視同股利）", "is.niCommon", cols=IS4); r += 1
qrow(r, "eps", "EPS（稀釋，美元）", "is.epsDil", scale=1, fmt="0.00", cols=IS4 + ["2026q2"]); r += 1
qrow(r, "shd", "稀釋加權股數（十億股）", "is.shDil", fmt="0.000", cols=IS4 + ["2026q2"]); r += 1
r += 1
Q.cell(r, 1, "現金流（9M；10-Q 現金流量表）").font = BOLD; r += 1
CF2 = ["9m2025", "9m2026"]
qrow(r, "cfo", "營運現金流", "cf.cfo", cols=CF2); r += 1
qrow(r, "cfda", "折舊攤銷", "cf.da", cols=CF2); r += 1
qrow(r, "cfsbc", "股權報酬", "cf.sbc", cols=CF2); r += 1
qrow(r, "cx", "資本支出（購置不動產與設備）", "cf.capex", neg=True, cols=CF2); r += 1
qrow(r, "chk_cx", "檢查：資本支出 − 分部資本支出合計", formula=lambda c: f"{c}{QR['cx']}-{c}{QR['capex_tot']}", fmt="0.000", tag="—", src="—", cols=CF2); r += 1
qrow(r, "ppefin", "以其他融資取得 PP&E（非現金）", "cf.ppeFin", cols=CF2); r += 1
qrow(r, "flp", "融資租賃本金", "cf.flPrin", neg=True, cols=CF2); r += 1
qrow(r, "fcf_co", "公司口徑 FCF（營運現金 − 資本支出）", formula=lambda c: f"{c}{QR['cfo']}-{c}{QR['cx']}", tag="Derived", src="—", cols=CF2); r += 1
r += 1
Q.cell(r, 1, "核對彙總（實績必須為 0）").font = BOLD; r += 1
QCHK = [k for k in QR if k.startswith("chk_")]
Q.cell(r, 1, "檢查：上列核對列絕對值合計"); QR["chk_all"] = r
Q[f"B{r}"] = "=" + "+".join(f"SUMPRODUCT(ABS(B{QR[k]}:I{QR[k]}))" for k in QCHK); Q[f"B{r}"].number_format = "0.000"; r += 1
Q.cell(r, 1, "檢查：9M26 − (2026Q1＋Q2＋Q3)（三分部營收、營業利益、調整後 EBITDA、資本支出）"); QR["chk_9m"] = r
Q[f"B{r}"] = "=" + "+".join(f"ABS(H{QR[k]}-D{QR[k]}-E{QR[k]}-F{QR[k]})" for k in ("space", "conn", "si", "ebit_space", "ebit_conn", "ebit_si", "adjEbitda_space", "adjEbitda_conn", "adjEbitda_si", "capex_space", "capex_conn", "capex_si"))
Q[f"B{r}"].number_format = "0.000"; r += 2

Q.cell(r, 1, "推導量（Seg、PL、Thesis 頁的起始值；[Derived]）").font = BOLD; r += 1
hdr(Q, r, ["項目", "數值", "說明"]); r += 1
def qd(key, label, formula, fmt="0.0%", note=""):
    global r
    Q.cell(r, 1, label); Q[f"B{r}"] = formula; Q[f"B{r}"].number_format = fmt; Q.cell(r, 3, note)
    dn(key, f"Q!$B${r}"); QR[key] = r; r += 1
qd("g0_ld", "9M 年增：Launch & Development", f"=H{QR['launchDev']}/G{QR['launchDev']}-1", note="9M26 ÷ 9M25 − 1（最近可用的 Verified 同期比較；近四季需 2025Q4，為二手）")
qd("g0_eg", "9M 年增：Enterprise & Government", f"=H{QR['entGov']}/G{QR['entGov']}-1")
qd("g0_ads", "9M 年增：Advertising（X）", f"=H{QR['ads']}/G{QR['ads']}-1")
qd("rpl0", "每次客戶發射收入（9M26）", f"=H{QR['launchSvc']}/H{QR['nlc']}", "0.0000", "US$B／次")
qd("m_sp0", "Space 調整後 EBITDA 率（9M26）", f"=H{QR['adjEbitda_space']}/H{QR['space']}")
qd("m_cn0", "Connectivity 調整後 EBITDA 率（9M26）", f"=H{QR['adjEbitda_conn']}/H{QR['conn']}")
qd("cx_sp0", "Space 資本支出 ÷ 營收（9M26）", f"=H{QR['capex_space']}/H{QR['space']}")
qd("cx_cn0", "Connectivity 資本支出 ÷ 營收（9M26）", f"=H{QR['capex_conn']}/H{QR['conn']}")
for s, lab in SEGS:
    qd(f"sbcr_{s}", f"{lab} 一般股權報酬 ÷ 營收（2026Q3）", f"=F{QR['sbc_' + s]}/F{QR[s]}", note="2026Q3 單季（IPO 後首個完整季）")
qd("da_sp0", "Space 年折舊（2026Q3 × 4）", f"=F{QR['da_space']}*4", "#,##0.000")
qd("da_cn0", "Connectivity 年折舊（2026Q3 × 4）", f"=F{QR['da_conn']}*4", "#,##0.000")
qd("uplift", "消費營收 ÷ (平均訂戶 × ARPU × 3)（2026Q3；含 Kits 等）", f"=F{QR['consumer']}/((E{QR['subs']}+F{QR['subs']})/2*F{QR['arpu']}/1000*3)", "0.000",
   "Kits（產品）收入全屬 Connectivity、未拆消費／企業（A 段第 9 節）→ 以本係數併入消費")
qd("avg_gw_q3", "2026Q3 平均在役名目算力", f"=(E{QR['gw']}+F{QR['gw']})/2", "0.000", "GW")
qd("ssi_q3", "2026Q3 SI Solutions & Infrastructure", f"=F{QR['ssi']}", "#,##0.000")
qd("ssi_h2", "2026Q2＋Q3 SI Solutions & Infrastructure（雲端服務自 Q2 開始）", f"=E{QR['ssi']}+F{QR['ssi']}", "#,##0.000")
qd("si_q3", "2026Q3 SI 營收", f"=F{QR['si']}", "#,##0.000")
qd("sie_q3", "2026Q3 SI 調整後 EBITDA", f"=F{QR['adjEbitda_si']}", "#,##0.000")
qd("ads_q3", "2026Q3 Advertising", f"=F{QR['ads']}", "#,##0.000")
qd("ld_9m", "9M26 Launch & Development", f"=H{QR['launchDev']}", "#,##0.000")
qd("eg_q3", "2026Q3 Enterprise & Government", f"=F{QR['entGov']}", "#,##0.000")
qd("subs_q3", "2026Q3 期末消費訂戶（百萬）", f"=F{QR['subs']}", "0.00")
for col in "BCDEFGHI": Q.column_dimensions[col].width = 11
Q.column_dimensions["A"].width = 64; Q.column_dimensions["J"].width = 22; Q.column_dimensions["K"].width = 90
Q.freeze_panes = "B4"
if "m_cn" not in OVR: IN[f"E{IN_ROW['m_cn']}"] = f"=Q!B{QR['m_cn0']}"
if "cx_cn" not in OVR: IN[f"E{IN_ROW['cx_cn']}"] = f"=Q!B{QR['cx_cn0']}"
def QA(k, p):
    return {"25": f"Q!I{QR[k]}", "m25": f"Q!G{QR[k]}", "1H": f"Q!H{QR[k]}"}[p]

# =====================================================================================================
# Thesis：命題模組（SI 名目算力 GW × 每 GW 年營收；GW → SI 投資 → 折舊；自用 GW → 影子效益）
TH["A1"] = ("Thesis — 命題模組：SI 名目算力（IT 口徑 GW；2026-09-30 1.7）由 SI 投資季速 ÷ Tokenomics 每 GW 資本支出推進；對外出租 GW（× ext_sh）× 每 GW 年營收（k × Tokenomics 持有成本，"
            "k 由 2026Q3 校準值收斂，情境軸，見 Seg 頁）→ SI 對外雲端服務子列；同一套 GW 推出 SI 投資、AI 折舊、營運成本；自用 GW → 讀法 2 影子效益")
TH["A2"] = ("三情境共用（情境軸是 k 與 Starlink 淨增，在 Seg 頁）。名目算力為 IT 口徑，與 Tokenomics Interface 第 1 列 IT 關鍵電力同口徑，不乘 PUE（審查第 3 條）；"
            "含對外出租、推定不含第三方雲（S6）。對外出租占比以 Customer B 衍生 72% 為起點，營收與 GW 同比例（審查第 4 條）")
T = {}
r = 4
TH.cell(r, 1, "A. 校準與起始值（單一數值，C 欄）").font = BOLD; r += 1
hdr(TH, r, ["項目", "單位", "數值"] + [""] * 10 + ["說明"]); r += 1
def tscal(key, label, unit, formula, fmt="0.000", note=""):
    global r
    TH.cell(r, 1, label); TH.cell(r, 2, unit); TH[f"C{r}"] = formula; TH[f"C{r}"].number_format = fmt; TH.cell(r, NOTE_C, note)
    dn(key, f"Thesis!$C${r}"); T[key] = r; r += 1
tscal("he0", "評價日機隊每 GW 年持有成本（Tokenomics IF_HoldEcon 依機隊占比加權）", "US$B/GW-年", "=mix0_h*he_h+mix0_b2*he_b2+mix0_b3*he_b3", "0.00")
tscal("opex0", "評價日機隊每 GW 年營運費用（IF_OpexGW 加權，不含折舊）", "US$B/GW-年", "=mix0_h*opx_h+mix0_b2*opx_b2+mix0_b3*opx_b3", "0.00")
tscal("rev_gw0", "2026Q3 每 GW 年營收＝SI S&I × 4 ÷ 平均在役 GW", "US$B/GW-年", "=ssi_q3*4/avg_gw_q3", "0.00",
      "審查第 4 條：2026Q3 年化 ÷ 平均在役 GW 校準；對外出租營收與 GW 同比例，故＝對外每 GW 年營收")
tscal("k0", "校準 k0＝2026Q3 每 GW 年營收 ÷ 評價日機隊持有成本", "x", "=rev_gw0/he0", "0.00", "S2 核心問題比值的起點（FY2028 比值見 Seg 頁各情境）")
tscal("rent_q3", "2026Q3 對外雲端服務營收（SI S&I × ext_sh）", "US$B", "=ext_sh*ssi_q3", "#,##0.000", "Customer B 衍生 5.028（事實總帳 cust.B.rev.2026q3）")
tscal("rent_9m", "9M26 對外雲端服務營收（(Q2＋Q3 SI S&I) × ext_sh；雲端服務自 Q2 開始）", "US$B", "=ext_sh*ssi_h2", "#,##0.000", "Customer B 9M 衍生 6.554")
tscal("nr_q3", "2026Q3 SI 其他收入（Grok／Cursor／X 訂閱、資料授權、API）", "US$B", "=ssi_q3-rent_q3", "#,##0.000", "含 Cursor 約 1.5 月 0.608")
tscal("m_so0", "SI 其他 EBITDA 率（2026Q3 校準：(SI 調整後 EBITDA − 對外出租 EBITDA) ÷ (SI 營收 − 對外出租)）", "%",
      "=(sie_q3-(rent_q3-ext_sh*avg_gw_q3*opex0/4))/(si_q3-rent_q3)", "0.0%", "對外出租 EBITDA＝營收 − 對外 GW × Tokenomics 營運費用；殘差含 Grok 訓練人力、評價日自用 GW 營運")
tscal("m_so", "SI 其他 EBITDA 率（預測期採用）", "%", "=m_so0+m_so_adj", "0.0%")
tscal("cxgw26", "每 GW 投資（2026 新增：GB300 錨 IT＋自建 × 機房）", "US$B/GW", "=cit_b3+selfb*cfac", "0.00", "Tokenomics IF_CapexIT｜GB300＋selfb × IF_CapexFacility")
tscal("ppe_gw", "對照：PP&E（伺服器與網通＋資料中心基礎設施，2026-09-30）÷ 名目算力", "US$B/GW",
      f"=({fb('ppe.servers.20260930')}+{fb('ppe.dcInfra.20260930')})/gw_0", "0.00", "毛額 43.206＋6.333 ÷ 1.7；含 Hopper 舊機隊、機房可能部分租用 → 低於 Tokenomics")
r += 1
TH.cell(r, 1, "B. SI 名目算力（GW，期末；IT 口徑）— 依 Tokenomics 世代錨分桶").font = BOLD; r += 1
hdr(TH, r, MODEL_HDR); r += 1
TH_HDR = r - 1
BK = [("h", "Hopper 錨"), ("b2", "GB200 錨"), ("b3", "GB300 錨"), ("vr", "VR200 錨")]
MIX0 = {"h": "mix0_h", "b2": "mix0_b2", "b3": "mix0_b3", "vr": "0"}
def addmix(p, b):
    y = PY[p]
    if y == 2026: return {"b3": "1"}.get(b, "0")
    if y == 2027: return {"b3": "am_27", "vr": "(1-am_27)"}.get(b, "0")
    return {"vr": "1"}.get(b, "0")
def tput(key, label, unit, vals, fmt="#,##0.000", note=None, **kw):
    global r
    put(TH, r, label, unit, vals, fmt, note, **kw); T[key] = r; r += 1
ADDS_R = r; END_R = r + 5
_v = {"2H": "=pace*0.25", 27: "=pace", 28: "=pace", "26": f"=F{r}"}
for p in (29, 30, 31, 32):
    _v[p] = f"=pace+(g_tv*{COL[prevp(p)]}{END_R}-pace)*MIN(1,({PY[p]}-2028)/conv_mw)"
tput("adds", "SI 名目算力新增（FY2029 起收斂到 g × 期初 GW）", "GW", _v, note="Q4 26 與 FY2027–28＝pace（SI 投資季速 × 4 ÷ 每 GW 投資）；FY2029–32 線性收斂（同 Alphabet 修正一 F5）")
for b, nm in BK:
    v = {"1H": f"=gw_0*{MIX0[b]}"}
    for p in FC:
        v[p] = f"={COL[prevp(p)]}{r}+{COL[p]}{ADDS_R}*{addmix(p, b)}"
    v["26"] = f"=F{r}"
    tput(f"end_{b}", f"期末 GW：{nm}", "GW", v, note="評價日 1.7 GW × 機隊占比（mix0）；新增：2026 GB300 錨、2027 GB300／VR200 各半、2028 起 VR200 錨" if b == "h" else None)
assert r == END_R
tput("end", "期末 SI 名目算力", "GW", {p: "=" + "+".join(f"{COL[p]}{T['end_' + b]}" for b, _ in BK) for p in ["1H"] + FC + ["26"]}, bold=True)
for b, nm in BK:
    v = {p: f"=({COL[prevp(p)]}{T['end_' + b]}+{COL[p]}{T['end_' + b]})/2" for p in FC}
    tput(f"avg_{b}", f"平均在役 GW：{nm}", "GW", v)
tput("avg", "平均在役 GW", "GW", {p: "=" + "+".join(f"{COL[p]}{T['avg_' + b]}" for b, _ in BK) for p in FC})
tput("avg_ext", "平均在役 GW：對外出租（× ext_sh）", "GW", {p: f"={COL[p]}{T['avg']}*ext_sh" for p in FC})
tput("avg_int", "平均在役 GW：自用（×(1 − ext_sh)）", "GW", {p: f"={COL[p]}{T['avg']}*(1-ext_sh)" for p in FC})
SUMHE = lambda c: "+".join(f"{c}{T['avg_' + b]}*he_{b}" for b, _ in BK)
SUMOP = lambda c: "+".join(f"{c}{T['avg_' + b]}*opx_{b}" for b, _ in BK)
tput("he_fl", "在役機隊每 GW 年持有成本（IF_HoldEcon 依平均在役 GW 加權）", "US$B/GW-年", {p: f"=({SUMHE(COL[p])})/{COL[p]}{T['avg']}" for p in FC}, "0.00",
     note="k × 本列＝每 GW 年營收（Seg 頁）；S2 核心問題的分母")
tput("opx_fl", "在役機隊每 GW 年營運費用（IF_OpexGW 加權）", "US$B/GW-年", {p: f"=({SUMOP(COL[p])})/{COL[p]}{T['avg']}" for p in FC}, "0.00")
tput("ext_opex", "對外出租營運費用（對外 GW × 營運費用；含電費）", "US$B", {**{p: f"={COL[p]}{T['avg_ext']}*{COL[p]}{T['opx_fl']}*{LEN[p]}" for p in FC}, "26": f"=F{r}"}, "#,##0.00")
SUMOP0 = "+".join(f"$E${T['end_' + b]}*opx_{b}" for b, _ in BK)
tput("int_opex", "自用增量營運費用（相對評價日自用機隊；Tokenomics IF_OpexGW）", "US$B",
     {**{p: f"=(1-ext_sh)*(({SUMOP(COL[p])})-({SUMOP0}))*{LEN[p]}" for p in FC}, "26": f"=F{r}"}, "#,##0.00", note="評價日自用機隊的營運費用已在 2026Q3 校準的 SI 其他 EBITDA 率內")
r += 1
TH.cell(r, 1, "C. SI 投資（新增 GW × Tokenomics 每 GW 成本；汰換只換 IT；含以其他融資取得之 PP&E）").font = BOLD; r += 1
SUMADD = lambda p: "+".join(f"{addmix(p, b)}*cit_{b}" for b, _ in BK if b in ("b3", "vr"))
tput("cx_it", "SI IT 投資（成長）", "US$B", {**{p: f"={COL[p]}{T['adds']}*({SUMADD(p)})" for p in FC}, "26": f"=F{r}"}, "#,##0.00")
tput("cx_fac", "SI 機房投資（自建比例 × IF_CapexFacility）", "US$B", {**{p: f"={COL[p]}{T['adds']}*selfb*cfac" for p in FC}, "26": f"=F{r}"}, "#,##0.00")
tput("retire", "汰換 GW（6 年前上線，換成 VR200 錨 IT）", "GW", {"2H": 0, 27: 0, 28: 0, 29: 0, 30: "=gw_ye24", 31: "=gw_ye25-gw_ye24", 32: f"=gw_0-gw_ye25+F{T['adds']}", "26": f"=F{r}"},
     note="IT 壽命 6 年：2024 年底在役 → 2030 汰換，以此類推；汰換後 GW 照常在役")
tput("cx_ref", "SI 汰換投資（只含 IT）", "US$B", {**{p: f"={COL[p]}{T['retire']}*cit_vr" for p in FC}, "26": f"=F{r}"}, "#,##0.00")
tput("si_inv", "SI 投資合計（現金資本支出＋以其他融資取得）", "US$B", {**{p: f"={COL[p]}{T['cx_it']}+{COL[p]}{T['cx_fac']}+{COL[p]}{T['cx_ref']}" for p in FC}, "1H": f"={QA('capex_si', '1H')}+ppefin_9m",
     "26": f"=E{r}+F{r}"}, "#,##0.00", bold=True, note="9M26＝SI 資本支出 39.995＋以其他融資取得 3.921（推定全屬 SI）")
tput("si_cash", "  其中：SI 現金資本支出（×(1 − fin_sh)）", "US$B", {**{p: f"={COL[p]}{T['si_inv']}*(1-fin_sh)" for p in FC}, "1H": "=" + QA("capex_si", "1H"), "26": f"=E{r}+F{r}"}, "#,##0.00")
tput("si_fin", "  其中：以其他融資取得 PP&E（× fin_sh；非現金）", "US$B", {**{p: f"={COL[p]}{T['si_inv']}*fin_sh" for p in FC}, "1H": "=ppefin_9m", "26": f"=E{r}+F{r}"}, "#,##0.00")
tput("inv_gw", "  對照：SI 投資 ÷ 新增 GW", "US$B/GW", {p: f"=IF({COL[p]}{T['adds']}>0,({COL[p]}{T['cx_it']}+{COL[p]}{T['cx_fac']})/{COL[p]}{T['adds']},0)" for p in FC}, "0.00")
r += 1
TH.cell(r, 1, "D. 新增 AI 折舊（在建工程與新增投資依年限直線攤提；評價日在役資產的折舊見 E 節基礎折舊）").font = BOLD; r += 1
tput("cip_it", "在建工程轉入（SI）：IT（2026 年度批次）", "US$B", {"2H": "=cip_0*cip_si*cip_it"}, "#,##0.00", note="同 Alphabet 修正三：在建工程為已付資產，照時程折舊、不另增 GW（判斷題）")
tput("cip_fac", "在建工程轉入（SI）：機房", "US$B", {"2H": "=cip_0*cip_si*(1-cip_it)"}, "#,##0.00")
tput("v_it", "IT 批次金額（成長＋汰換）", "US$B", {p: f"={COL[p]}{T['cx_it']}+{COL[p]}{T['cx_ref']}" for p in FC}, "#,##0.00")
tput("v_fac", "機房批次金額", "US$B", {p: f"={COL[p]}{T['cx_fac']}" for p in FC}, "#,##0.00")
def depf(amt, age, life):
    a = f"({age})"
    return f"{amt}/{life}*IF({a}<0,0,IF({a}=0,phi_0,IF({a}<{life},1,IF({a}={life},1-phi_0,0))))"
def dep_row(srcs, life, lenq=True):
    out = {}
    for p in FC:
        terms = [depf(a, PY[p] - vy, life) for a, vy in srcs if PY[p] >= vy]
        out[p] = "=(" + "+".join(terms) + ")" + ("*0.25" if (p == "2H" and lenq) else "")
    return out
it_src = [(f"$F${T['cip_it']}", 2026)] + [(f"${COL[p]}${T['v_it']}", PY[p]) for p in FC]
fac_src = [(f"$F${T['cip_fac']}", 2026)] + [(f"${COL[p]}${T['v_fac']}", PY[p]) for p in FC]
d = dep_row(it_src, "life_it"); d["26"] = f"=F{r}"; tput("da_it", "新增 AI 折舊：IT", "US$B", d, "#,##0.00", note="Q4 26 欄＝年折舊 × 0.25")
d = dep_row(fac_src, "life_fac"); d["26"] = f"=F{r}"; tput("da_fac", "新增 AI 折舊：機房", "US$B", d, "#,##0.00")
tput("da_ai", "新增 AI 折舊合計", "US$B", {p: f"={COL[p]}{T['da_it']}+{COL[p]}{T['da_fac']}" for p in FC + ["26"]}, "#,##0.00", bold=True)
tput("da_ext", "  其中：對外（× ext_sh）", "US$B", {p: f"={COL[p]}{T['da_ai']}*ext_sh" for p in FC + ["26"]}, "#,##0.00")
tput("da_int", "  其中：自用（×(1 − ext_sh)）", "US$B", {p: f"={COL[p]}{T['da_ai']}*(1-ext_sh)" for p in FC + ["26"]}, "#,##0.00")
r += 1
TH.cell(r, 1, "E. SI 評價日在役資產的基礎折舊（隨期初機隊汰換遞減）").font = BOLD; r += 1
_ret = T["retire"]
_cv = {"2H": f"=mid*F{_ret}"}
for p in ANN:
    prev_cols = [COL[q] for q in FC if PY[q] < PY[p]]
    _cv[p] = "=" + "+".join(f"{c}{_ret}" for c in prev_cols) + f"+mid*{COL[p]}{_ret}"
_cv["26"] = f"=F{r}"
tput("cumret", "累計已汰換的期初機隊 GW（當年以半年計）", "GW", _cv)
tput("dpi0", "期初機隊每 GW IT 年折舊（Tokenomics IF_DeprIT 依期初占比加權）", "US$B/GW-年", {"1H": "=mix0_h*dpi_h+mix0_b2*dpi_b2+mix0_b3*dpi_b3"}, "0.000")
tput("da_b", "SI 基礎折舊＝da_si0 − MIN(da_si0 × IT 占比, 累計汰換 GW × 每 GW IT 折舊)", "US$B",
     {**{p: f"=(da_si0-MIN(da_si0*it_sh_da,{COL[p]}{T['cumret']}*$E${T['dpi0']}))*{LEN[p]}" for p in FC}, "26": f"=F{r}"}, "#,##0.00",
     note="2026Q3 SI 折舊攤銷 2.473 × 4（含 X 無形資產攤銷）；期初機隊汰換後由新購 IT 折舊取代")
r += 1
TH.cell(r, 1, "F. 自用 AI（讀法 2 影子效益：自用增量 GW × Tokenomics 持有成本，k＝1；非財報營收）").font = BOLD; r += 1
SUMHE0 = "(" + "+".join(f"$E${T['end_' + b]}*he_{b}" for b, _ in BK) + ")"
tput("shadow", "自用 AI 影子效益（自用增量 GW × IF_HoldEcon）", "US$B",
     {**{p: f"=(1-ext_sh)*(({SUMHE(COL[p])})-{SUMHE0})*{LEN[p]}" for p in FC}, "26": f"=F{r}"}, "#,##0.00",
     note="原則第 7 節讀法 2：只進評價口徑（PL「營業利益（評價口徑）」、Val）；營收列與分部對帳不動")
tput("int_cost", "自用 AI 增量成本（營運＋自用新增折舊）", "US$B", {p: f"={COL[p]}{T['int_opex']}+{COL[p]}{T['da_int']}" for p in FC + ["26"]}, "#,##0.00")
tput("sh_vs", "  影子效益 ÷ 自用 AI 增量成本", "x", {p: f"=IF({COL[p]}{T['int_cost']}>0,{COL[p]}{T['shadow']}/{COL[p]}{T['int_cost']},0)" for p in FC}, "0.00",
     note="接近 1＝價值中性（持有成本以 10% 資金成本回收資本）")
widths(TH)
TH.freeze_panes = "C4"

# =====================================================================================================
# Seg：營收分部層（三情境各一完整區塊）
SG["A1"] = ("Seg — 營收分部層：官方分部 Space（Launch Services、Launch & Development）、Connectivity（Consumer、Enterprise & Government）、SI（Advertising、SI Solutions & Infrastructure）；"
            "SI S&I 之下的命題子列＝SI 對外雲端服務＋Grok／Cursor／X 訂閱、資料授權與 API（子列合計＝SI S&I）。情境軸：SI 每 GW 年營收的 k（FY2030 到達）與 Starlink 消費訂戶淨增")
SG["A2"] = ("成長路徑：起始年增＝9M26 對 9M25 年增（Q 頁；近四季需 2025Q4，只有二手），FY2030 收斂到長期值（conv_n）。Space＝客戶發射 × 每次收入＋發射與開發；"
            "Connectivity＝平均消費訂戶 × ARPU × 3 × 係數（含 Kits）＋企業與政府；SI 對外出租＝對外 GW × k × 持有成本（Thesis）；SI 其他以 Q4 26 備考全季（含 Cursor 全季）為基期")
SCN = {"A": ("k30_A", "nadd_A", "情境 A（算力租金維持高檔、Starlink 增長快）"), "B": ("k30_B", "nadd_B", "情境 B（租金向持有成本部分收斂；基準）"),
       "C": ("k30_C", "nadd_C", "情境 C（租金收斂到持有成本、Starlink 增長放緩）")}
SR = {s: {} for s in "ABC"}
r = 4
def gpath(g0, gl):
    return {**{"2H": f"={g0}"}, **{p: f"={g0}+({gl}-{g0})*MIN(1,({PY[p]}-2026)/conv_n)" for p in ANN}}
def sput(s, key, label, unit, vals, fmt="#,##0.000", note=None, **kw):
    global r
    put(SG, r, label, unit, vals, fmt, note, **kw); SR[s][key] = r; r += 1
for s in "ABC":
    kk, nn, title = SCN[s]
    SG.cell(r, 1, title).font = BOLD; SG.cell(r, 1).fill = SUB; r += 1
    hdr(SG, r, MODEL_HDR); r += 1
    R = SR[s]
    # --- SI ---
    sput(s, "k", "  k＝每 GW 年營收 ÷ 持有成本（Q4 26＝校準 k0，FY2030＝情境值）", "x", {**{"2H": "=k0"}, **{p: f"=k0+({kk}-k0)*MIN(1,({PY[p]}-2026)/conv_k)" for p in ANN}}, "0.00")
    sput(s, "rev_gw", "  對外每 GW 年營收＝k × 在役機隊持有成本", "US$B/GW-年", {p: f"={COL[p]}{R['k']}*Thesis!{COL[p]}{T['he_fl']}" for p in FC}, "0.00")
    v = {"1H": "=Thesis!C{0}".format(T["rent_9m"])}
    for p in FC: v[p] = f"=Thesis!{COL[p]}{T['avg_ext']}*{COL[p]}{R['rev_gw']}*{LEN[p]}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "rent", "SI 對外雲端服務（命題子列）", "US$B", v, note="對外 GW × 每 GW 年營收；9M26＝(Q2＋Q3 SI S&I) × ext_sh（Customer B 衍生）")
    sput(s, "g_nr", "  SI 其他收入年增", "%", {**{p: f"=g27_nr+(gl_nr-g27_nr)*MIN(1,({PY[p]}-2027)/(conv_n-1))" for p in ANN}}, "0.0%", "FY2027＝g27_nr（以 Q4 26 備考年化為基期），FY2030 收斂到 gl_nr")
    v = {"1H": f"={QA('ssi', '1H')}-E{R['rent']}", "2H": "=nr_q3+cur_pf", 27: f"=F{r}*4*(1+H{R['g_nr']})^0.625", "26": f"=E{r}+F{r}"}
    for p in ANN[1:]: v[p] = f"={COL[prevfy(p)]}{r}*(1+{COL[p]}{R['g_nr']})"
    sput(s, "nr", "Grok／Cursor／X 訂閱、資料授權與 API（命題子列）", "US$B", v, note="Q4 26＝2026Q3 值＋Cursor 備考補足（審查第 8 條：以備考全季為起點）；FY2027＝Q4 26 × 4 ×(1＋年增)^0.625")
    v = {p: "=" + QA("ssi", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"={COL[p]}{R['rent']}+{COL[p]}{R['nr']}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "ssi", "SI Solutions & Infrastructure", "US$B", v, bold=True)
    sput(s, "g_ads", "  Advertising 年增", "%", gpath("g0_ads", "gl_ads"), "0.0%", "起點＝9M 年增（Q 頁），FY2030 收斂到 gl_ads")
    v = {p: "=" + QA("ads", p) for p in ("m25", "1H")}
    v["2H"] = f"=ads_q3*(1+F{R['g_ads']})^0.25"; v["26"] = f"=E{r}+F{r}"
    for p in ANN: v[p] = f"={COL[prevfy(p)]}{r}*(1+{COL[p]}{R['g_ads']})"
    sput(s, "ads", "Advertising", "US$B", v, bold=True)
    v = {p: "=" + QA("si", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{R['ads']}+{COL[p]}{R['ssi']}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "si", "SI (Super Intelligence; formerly AI)", "US$B", v, bold=True, fill=SUB)
    # --- Space ---
    sput(s, "nl", "  客戶發射次數", "次", {"1H": "=" + QA("nlc", "1H"), "2H": "=nl_q4", 27: "=nl_27", **{p: f"={COL[prevfy(p)]}{r}*(1+g_nl)" for p in ANN[1:]}, "26": f"=E{r}+F{r}"}, "0.0")
    sput(s, "rpl", "  每次客戶發射收入", "US$B／次", {"1H": "=rpl0", "2H": "=rpl0", **{p: f"=rpl0*(1+g_rpl)^({PY[p]}-2026)" for p in ANN}}, "0.0000")
    v = {p: "=" + QA("launchSvc", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"={COL[p]}{R['nl']}*{COL[p]}{R['rpl']}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "lsv", "Launch Services", "US$B", v, bold=True)
    sput(s, "g_ld", "  Launch & Development 年增", "%", gpath("g0_ld", "gl_ld"), "0.0%")
    v = {p: "=" + QA("launchDev", p) for p in ("m25", "1H")}
    v["2H"] = f"=ld_9m/3*(1+F{R['g_ld']})^0.25"; v["26"] = f"=E{r}+F{r}"
    for p in ANN: v[p] = f"={COL[prevfy(p)]}{r}*(1+{COL[p]}{R['g_ld']})"
    sput(s, "ldv", "Launch & Development", "US$B", v, bold=True, note="Q4 26＝9M26 季均 ×(1＋年增)^0.25（里程碑收入不規則）")
    v = {p: "=" + QA("space", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{R['lsv']}+{COL[p]}{R['ldv']}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "space", "Space", "US$B", v, bold=True, fill=SUB)
    # --- Connectivity ---
    v = {"1H": "=subs_q3", "2H": "=E{0}+nadd_q4".format(r), 27: f"=F{r}+{nn}"}
    for p in ANN[1:]: v[p] = f"={COL[prevp(p)]}{r}+{nn}*(1-nadd_dec)^({PY[p]}-2027)"
    sput(s, "subs", "  Starlink 消費訂戶（期末，百萬）", "百萬", v, "0.00")
    sput(s, "subs_avg", "  平均消費訂戶", "百萬", {p: f"=({COL[prevp(p)]}{R['subs']}+{COL[p]}{R['subs']})/2" for p in FC}, "0.00")
    sput(s, "arpu", "  消費 ARPU（美元／月）", "US$", {"2H": "=arpu0*(1+g_arpu)^0.25", **{p: f"=arpu0*(1+g_arpu)^({PY[p]}+0.5-2026.75)" for p in ANN}}, "0.00")
    v = {p: "=" + QA("consumer", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"={COL[p]}{R['subs_avg']}*{COL[p]}{R['arpu']}/1000*12*{LEN[p]}*uplift"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "cons", "Consumer", "US$B", v, bold=True, note="平均訂戶 × ARPU × 月數 × uplift（2026Q3 校準，含 Kits）")
    sput(s, "g_eg", "  Enterprise & Government 年增", "%", gpath("g0_eg+g0_eg_adj", "gl_eg"), "0.0%", "起點＝9M 年增 +70%（Q 頁），FY2030 收斂到 gl_eg")
    v = {p: "=" + QA("entGov", p) for p in ("m25", "1H")}
    v["2H"] = f"=eg_q3*(1+F{R['g_eg']})^0.25"; v["26"] = f"=E{r}+F{r}"
    for p in ANN: v[p] = f"={COL[prevfy(p)]}{r}*(1+{COL[p]}{R['g_eg']})"
    sput(s, "eg", "Enterprise & Government", "US$B", v, bold=True)
    v = {p: "=" + QA("conn", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{R['cons']}+{COL[p]}{R['eg']}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "conn", "Connectivity", "US$B", v, bold=True, fill=SUB)
    v = {p: "=" + QA("rev", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{R['space']}+{COL[p]}{R['conn']}+{COL[p]}{R['si']}"
    v["26"] = f"=E{r}+F{r}"
    sput(s, "rev", "合併營收", "US$B", v, bold=True)
    P2 = ["m25", "1H"] + FC + ["26"]
    sput(s, "chk1", "檢查：Space 子列合計 − Space", "US$B", {p: f"={COL[p]}{R['lsv']}+{COL[p]}{R['ldv']}-{COL[p]}{R['space']}" for p in P2}, "0.000")
    sput(s, "chk2", "檢查：Connectivity 子列合計 − Connectivity", "US$B", {p: f"={COL[p]}{R['cons']}+{COL[p]}{R['eg']}-{COL[p]}{R['conn']}" for p in P2}, "0.000")
    sput(s, "chk3", "檢查：SI 子列合計 − SI", "US$B", {p: f"={COL[p]}{R['ads']}+{COL[p]}{R['ssi']}-{COL[p]}{R['si']}" for p in P2}, "0.000")
    sput(s, "chk4", "檢查：SI S&I 命題子列合計 − SI Solutions & Infrastructure", "US$B", {p: f"={COL[p]}{R['rent']}+{COL[p]}{R['nr']}-{COL[p]}{R['ssi']}" for p in ["1H"] + FC + ["26"]}, "0.000")
    sput(s, "chk5", "檢查：分部合計 − 合併營收", "US$B", {p: f"={COL[p]}{R['space']}+{COL[p]}{R['conn']}+{COL[p]}{R['si']}-{COL[p]}{R['rev']}" for p in YK}, "0.000")
    sput(s, "chk6", "檢查：實績欄合併營收 − Q 頁", "US$B", {p: f"={COL[p]}{R['rev']}-{QA('rev', p)}" for p in HIST}, "0.000")
    sput(s, "s2", "S2 核心問題：SI 每 GW 年營收 ÷ Tokenomics 每 GW 年持有成本（＝k）", "x", {p: f"={COL[p]}{R['k']}" for p in FC}, "0.00",
         note="工作單 S2：與 OpenAI PxQ 頁同口徑；FY2028 欄＝k 假設路徑（k0 → 情境值，conv_k 年）的內插值，不是模型產出（> 1＝租金高於打平持有成本）")
    # 非 AI 資本支出與折舊
    sput(s, "cxp_sp", "  Space 資本支出占營收", "%", {**{"2H": "=cx_sp0"}, **{p: f"=cx_sp0+(cx_spl-cx_sp0)*MIN(1,({PY[p]}-2026)/conv_n)" for p in ANN}}, "0.0%")
    sput(s, "cx_sp", "資本支出：Space", "US$B", {**{p: f"={COL[p]}{R['space']}*{COL[p]}{R['cxp_sp']}" for p in FC}, "26": f"=F{r}"})
    sput(s, "cx_cn", "資本支出：Connectivity（營收 × cx_cn）", "US$B", {**{p: f"={COL[p]}{R['conn']}*cx_cn" for p in FC}, "26": f"=F{r}"})
    d = dep_row([(f"${COL[p]}${R['cx_sp']}", PY[p]) for p in FC], "life_sp"); d["26"] = f"=F{r}"
    sput(s, "dn_sp", "新增折舊：Space（年限 life_sp）", "US$B", d)
    d = dep_row([(f"${COL[p]}${R['cx_cn']}", PY[p]) for p in FC], "life_cn"); d["26"] = f"=F{r}"
    sput(s, "dn_cn", "新增折舊：Connectivity（年限 life_cn）", "US$B", d)
    r += 1
widths(SG)
SG.freeze_panes = "C4"

# =====================================================================================================
# PL：分部營收 → 分部調整後 EBITDA − 折舊 − 股權報酬 → 分部營業利益 → 合併 → 業外 → 稅 → EPS（三情境）
PL["A1"] = ("PL — 分部營收 → 分部調整後 EBITDA（公司定義，不含 SBC）− 折舊攤銷 − 股權報酬（含 Cursor 收購後服務期獎酬）→ 分部營業利益 → 合併 → 利息 → 稅 → 淨利 → EPS（三情境；US$B，EPS 美元）。"
            "實績欄 FY2025*（二手轉述 S-1，只有分部營收與營業利益）、9M25A、9M26A 取自 Q 頁；FY2026E＝9M 實績＋Q4 模型")
PL["A2"] = ("分部調整後 EBITDA：Space＝營收 × 利益率（9M26 −5.9% → FY2030 m_spl）；Connectivity＝營收 × m_cn（9M26）；SI＝對外出租營收 − 對外 GW 營運費用（Tokenomics）＋SI 其他收入 × m_so（2026Q3 校準）"
            " − 自用增量營運費用。折舊：Space／Connectivity 基礎折舊（2026Q3 × 4，直線遞減到 0）＋新增資本支出折舊；SI＝基礎折舊（汰換遞減）＋新增 AI 折舊。三分部合計＝合併（無公司層級調節項）")
PR_ = {s: {} for s in "ABC"}
r = 4
PL.cell(r, 1, "分部對帳彙總（三情境 × 全部年度取絕對值最大者；實績欄同列）").font = BOLD; r += 1
hdr(PL, r, MODEL_HDR); r += 1
SUMROW = r; r += 7
def pput(s, key, label, unit, vals, fmt="#,##0.000", note=None, **kw):
    global r
    put(PL, r, label, unit, vals, fmt, note, **kw); PR_[s][key] = r; r += 1
def tmid(p): return "2026.875" if p == "2H" else f"{PY[p]}.5"
for s in "ABC":
    R = SR[s]; P = PR_[s]
    PL.cell(r, 1, SCN[s][2]).font = BOLD; PL.cell(r, 1).fill = SUB; r += 1
    hdr(PL, r, MODEL_HDR); r += 1
    ref = lambda k: {p: f"=Seg!{COL[p]}{R[k]}" for p in YK if SG[f"{COL[p]}{R[k]}"].value is not None}
    pput(s, "lsv", "Launch Services", "US$B", ref("lsv"))
    pput(s, "ldv", "Launch & Development", "US$B", ref("ldv"))
    pput(s, "space", "Space", "US$B", ref("space"), bold=True, fill=SUB)
    pput(s, "cons", "Consumer", "US$B", ref("cons"))
    pput(s, "eg", "Enterprise & Government", "US$B", ref("eg"))
    pput(s, "conn", "Connectivity", "US$B", ref("conn"), bold=True, fill=SUB)
    pput(s, "ads", "Advertising", "US$B", ref("ads"))
    pput(s, "rent", "SI 對外雲端服務（命題子列）", "US$B", ref("rent"))
    pput(s, "nr", "Grok／Cursor／X 訂閱、資料授權與 API（命題子列）", "US$B", ref("nr"))
    pput(s, "ssi", "SI Solutions & Infrastructure", "US$B", ref("ssi"))
    pput(s, "si", "SI (Super Intelligence; formerly AI)", "US$B", ref("si"), bold=True, fill=SUB)
    pput(s, "rev", "合併營收", "US$B", ref("rev"), bold=True)
    pput(s, "g", "  合併營收年增", "%", {p: f"={COL[p]}{P['rev']}/{COL[q]}{P['rev']}-1" for p, q in (("1H", "m25"), ("26", "25"), (27, "26"), (28, 27), (29, 28), (30, 29), (31, 30), (32, 31))}, "0.0%")
    P2 = ["m25", "1H"] + FC + ["26"]
    pput(s, "chk1", "檢查：Space 子列合計 − Space", "US$B", {p: f"={COL[p]}{P['lsv']}+{COL[p]}{P['ldv']}-{COL[p]}{P['space']}" for p in P2}, "0.000")
    pput(s, "chk2", "檢查：Connectivity 子列合計 − Connectivity", "US$B", {p: f"={COL[p]}{P['cons']}+{COL[p]}{P['eg']}-{COL[p]}{P['conn']}" for p in P2}, "0.000")
    pput(s, "chk3", "檢查：SI 子列合計 − SI", "US$B", {p: f"={COL[p]}{P['ads']}+{COL[p]}{P['ssi']}-{COL[p]}{P['si']}" for p in P2}, "0.000")
    pput(s, "chk3b", "檢查：SI S&I 命題子列合計 − SI Solutions & Infrastructure", "US$B", {p: f"={COL[p]}{P['rent']}+{COL[p]}{P['nr']}-{COL[p]}{P['ssi']}" for p in ["1H"] + FC + ["26"]}, "0.000")
    pput(s, "chk4", "檢查：分部合計 − 合併營收", "US$B", {p: f"={COL[p]}{P['space']}+{COL[p]}{P['conn']}+{COL[p]}{P['si']}-{COL[p]}{P['rev']}" for p in YK}, "0.000")
    # 分部調整後 EBITDA
    pput(s, "msp", "  Space 調整後 EBITDA 率（9M26 → FY2030 m_spl）", "%", {**{"2H": "=m_sp0"}, **{p: f"=m_sp0+(m_spl-m_sp0)*MIN(1,({PY[p]}-2026)/conv_n)" for p in ANN}}, "0.0%")
    v = {p: "=" + QA("adjEbitda_space", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{P['space']}*{COL[p]}{P['msp']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "e_sp", "Space 分部調整後 EBITDA", "US$B", v)
    v = {p: "=" + QA("adjEbitda_conn", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{P['conn']}*m_cn"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "e_cn", "Connectivity 分部調整後 EBITDA", "US$B", v)
    pput(s, "e_rent", "  SI 對外出租 EBITDA（營收 − 對外 GW × Tokenomics 營運費用）", "US$B", {**{p: f"={COL[p]}{P['rent']}-Thesis!{COL[p]}{T['ext_opex']}" for p in FC}, "26": f"=F{r}"})
    pput(s, "e_oth", "  SI 其他 EBITDA（(SI 其他收入＋廣告) × m_so − 自用增量營運費用）", "US$B",
         {**{p: f"=({COL[p]}{P['nr']}+{COL[p]}{P['ads']})*m_so-Thesis!{COL[p]}{T['int_opex']}" for p in FC}, "26": f"=F{r}"})
    v = {p: "=" + QA("adjEbitda_si", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{P['e_rent']}+{COL[p]}{P['e_oth']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "e_si", "SI 分部調整後 EBITDA", "US$B", v)
    pput(s, "ebitda", "合併調整後 EBITDA（分部合計）", "US$B", {p: f"={COL[p]}{P['e_sp']}+{COL[p]}{P['e_cn']}+{COL[p]}{P['e_si']}" for p in YK})
    # 折舊攤銷
    v = {p: "=" + QA("da_space", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"=da_sp0*{LEN[p]}*MAX(0,1-({tmid(p)}-t_val)/rem_sp)+Seg!{COL[p]}{R['dn_sp']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "d_sp", "Space 折舊攤銷（基礎遞減＋新增）", "US$B", v)
    v = {p: "=" + QA("da_conn", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"=da_cn0*{LEN[p]}*MAX(0,1-({tmid(p)}-t_val)/rem_cn)+Seg!{COL[p]}{R['dn_cn']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "d_cn", "Connectivity 折舊攤銷（基礎遞減＋新增）", "US$B", v)
    v = {p: "=" + QA("da_si", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"=Thesis!{COL[p]}{T['da_b']}+Thesis!{COL[p]}{T['da_ai']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "d_si", "SI 折舊攤銷（基礎＋新增 AI）", "US$B", v)
    pput(s, "da", "折舊攤銷合計", "US$B", {p: f"={COL[p]}{P['d_sp']}+{COL[p]}{P['d_cn']}+{COL[p]}{P['d_si']}" for p in ["m25", "1H"] + FC + ["26"]})
    pput(s, "da_ai", "  其中：新增 AI 折舊", "US$B", {**{p: f"=Thesis!{COL[p]}{T['da_ai']}" for p in FC}, "26": f"=F{r}"})
    # 股權報酬
    for sg, key in (("space", "s_sp"), ("conn", "s_cn"), ("si", "s_si")):
        v = {p: "=" + QA("sbc_" + sg, p) for p in ("m25", "1H")}
        for p in FC: v[p] = f"={COL[p]}{P[sg]}*sbcr_{sg}"
        v["26"] = f"=E{r}+F{r}"
        pput(s, key, {"space": "Space", "conn": "Connectivity", "si": "SI"}[sg] + " 股權報酬（一般；營收 × 2026Q3 比率）", "US$B", v)
    v = {"m25": 0, "1H": f"={QA('sbcAcq_si', '1H')}+{QA('acq_si', '1H')}+{QA('restr_si', '1H')}+{QA('imp_sp', '1H')}",
         "2H": "=sbc_cur_left/sbc_cur_yrs*0.25", 27: "=sbc_cur_left/sbc_cur_yrs", 28: "=sbc_cur_left/sbc_cur_yrs", 29: "=MAX(0,sbc_cur_left-2.25*sbc_cur_left/sbc_cur_yrs)", 30: 0, 31: 0, 32: 0, "26": f"=E{r}+F{r}"}
    pput(s, "s_acq", "SI 收購相關股權報酬（Cursor）與其他調節項（收購成本、重組、減損）", "US$B", v,
         note="預測期只有 Cursor 收購後服務期獎酬（直線攤到 2029-06）；9M25 的重組與減損另含在下方實績差額")
    pput(s, "sbc", "股權報酬與其他調節合計", "US$B", {p: f"={COL[p]}{P['s_sp']}+{COL[p]}{P['s_cn']}+{COL[p]}{P['s_si']}+{COL[p]}{P['s_acq']}" for p in ["m25", "1H"] + FC + ["26"]})
    # 分部營業利益
    for sg, key, e, d_, s_, extra in (("space", "oi_sp", "e_sp", "d_sp", "s_sp", ""), ("conn", "oi_cn", "e_cn", "d_cn", "s_cn", ""), ("si", "oi_si", "e_si", "d_si", "s_si", "s_acq")):
        v = {p: "=" + QA("ebit_" + sg, p) for p in HIST}
        for p in FC: v[p] = f"={COL[p]}{P[e]}-{COL[p]}{P[d_]}-{COL[p]}{P[s_]}" + (f"-{COL[p]}{P[extra]}" if extra else "")
        v["26"] = f"=E{r}+F{r}"
        pput(s, key, {"space": "Space", "conn": "Connectivity", "si": "SI"}[sg] + " 分部營業利益", "US$B", v, bold=True)
    pput(s, "oi_seg", "分部營業利益合計", "US$B", {p: f"={COL[p]}{P['oi_sp']}+{COL[p]}{P['oi_cn']}+{COL[p]}{P['oi_si']}" for p in YK})
    v = {p: "=" + QA("ebit_tot", p) for p in HIST}
    for p in FC: v[p] = f"={COL[p]}{P['ebitda']}-{COL[p]}{P['da']}-{COL[p]}{P['sbc']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "oi", "合併營業利益", "US$B", v, bold=True, fill=SUB, note="預測期＝合併調整後 EBITDA − 折舊攤銷 − 股權報酬與調節（與分部合計獨立計算）")
    pput(s, "chk5", "檢查：分部營業利益合計 − 合併營業利益", "US$B", {p: f"={COL[p]}{P['oi_seg']}-{COL[p]}{P['oi']}" for p in YK}, "0.000")
    pput(s, "shadow", "自用 AI 影子效益（讀法 2；非財報營收）", "US$B", {**{p: f"=Thesis!{COL[p]}{T['shadow']}" for p in FC}, "26": f"=F{r}"})
    pput(s, "oiv", "營業利益（評價口徑；讀法 2 含影子效益）", "US$B", {**{p: f"={COL[p]}{P['oi']}+(internal_ai_reading=2)*{COL[p]}{P['shadow']}" for p in FC + ["26"]},
                                                         **{p: f"={COL[p]}{P['oi']}" for p in HIST}}, bold=True)
    pput(s, "om", "  營業利益率", "%", {p: f"={COL[p]}{P['oi']}/{COL[p]}{P['rev']}" for p in YK}, "0.0%")
    pput(s, "om_si", "  SI 營業利益率", "%", {p: f"={COL[p]}{P['oi_si']}/{COL[p]}{P['si']}" for p in YK}, "0.0%")
    pput(s, "em", "  調整後 EBITDA 率", "%", {p: f"={COL[p]}{P['ebitda']}/{COL[p]}{P['rev']}" for p in YK}, "0.0%")
    P["ii"] = r; r += 1; P["ie"] = r; r += 1
    v = {p: f"={QA('oie', p)}" for p in ("m25", "1H")}
    for p in FC: v[p] = 0
    v["26"] = f"=E{r}+F{r}"
    pput(s, "oie", "其他收益（預測期 0）", "US$B", v, note="9M26 −0.937（含債務清償損失 1.544、TBC 利益 0.561）")
    v = {p: f"={COL[p]}{P['oi']}+{QA('intExp', p)}+{QA('intInc', p)}+{COL[p]}{P['oie']}" for p in ("m25", "1H")}
    for p in FC: v[p] = f"={COL[p]}{P['oi']}+{COL[p]}{P['ii']}-{COL[p]}{P['ie']}+{COL[p]}{P['oie']}"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "pt", "稅前利益", "US$B", v)
    v = {p: "=" + QA("taxa", p) for p in ("m25", "1H")}
    for p in FC: v[p] = f"={COL[p]}{P['pt']}*tax"
    v["26"] = f"=E{r}+F{r}"
    pput(s, "tx", "所得稅", "US$B", v)
    pput(s, "ni", "淨利", "US$B", {p: f"={COL[p]}{P['pt']}-{COL[p]}{P['tx']}" for p in ["m25", "1H"] + FC + ["26"]}, bold=True)
    pput(s, "chk6", "檢查：實績淨利 − Q 頁", "US$B", {p: f"={COL[p]}{P['ni']}-{QA('ni', p)}" for p in ("m25", "1H")}, "0.000")
    pput(s, "dd", "視同股利（特別股轉換）", "US$B", {"1H": "=deemed_1H", **{p: 0 for p in FC}, "26": f"=E{r}+F{r}"})
    pput(s, "nic", "歸屬普通股淨利", "US$B", {p: f"={COL[p]}{P['ni']}-{COL[p]}{P['dd']}" for p in ["1H"] + FC + ["26"]})
    P["she"] = r; r += 1; P["sh"] = r; r += 1
    pput(s, "eps", "EPS（GAAP 稀釋，美元）", "US$", {p: f"={COL[p]}{P['nic']}/{COL[p]}{P['sh']}" for p in ["1H"] + FC + ["26"]}, "0.00", bold=True)
    pput(s, "epsv", "EPS（評價口徑；讀法 2 含稅後影子效益，美元）", "US$", {p: f"=({COL[p]}{P['nic']}+(internal_ai_reading=2)*{COL[p]}{P['shadow']}*(1-tax))/{COL[p]}{P['sh']}" for p in FC + ["26"]}, "0.00")
    r += 1
widths(PL)
PL.freeze_panes = "C4"
CHKS = [("chk1", "檢查：Space 子列合計 − Space"), ("chk2", "檢查：Connectivity 子列合計 − Connectivity"), ("chk3", "檢查：SI 子列合計 − SI"),
        ("chk3b", "檢查：SI S&I 命題子列合計 − SI Solutions & Infrastructure"), ("chk4", "檢查：分部合計 − 合併營收"), ("chk5", "檢查：分部營業利益合計 − 合併營業利益")]
for i, (k, lab) in enumerate(CHKS):
    rr = SUMROW + i
    PL.cell(rr, 1, lab); PL.cell(rr, 2, "US$B")
    for p in YK:
        c = COL[p]
        if PL[f"{c}{PR_['A'][k]}"].value is None: continue
        PL[f"{c}{rr}"] = "=MAX(" + ",".join(f"ABS({c}{PR_[s][k]})" for s in "ABC") + ")"
        PL[f"{c}{rr}"].number_format = "0.000"
    PL.cell(rr, NOTE_C, "三情境各區塊同名核對列的最大絕對值；必須為 0（tools/segment_check 讀本列）")

# =====================================================================================================
# CF：營運現金 − 資本支出 → 自由現金流（兩種口徑）；融資；收支關卡
CF["A1"] = ("CF — 現金流（US$B）。營運現金＝淨利＋折舊＋股權報酬＋營運資金；資本支出＝Space＋Connectivity（Seg）＋SI 現金資本支出（Thesis）。"
            "收支關卡口徑（審查第 2 條，10-30 更正）＝營運現金 − 資本支出 − 融資租賃本金 − 以其他融資取得之 PP&E（Valor 失敗售後租回本金不扣，列融資來源）；公司口徑 FCF＝營運現金 − 資本支出（並列）")
CF["A2"] = ("2026 年 9M：IPO 淨額 85.7、IPO 前發行 8.3、發債 51.8／償還 39.8（過橋與 xAI 舊債清償）。預測期不配息、不買回；EchoStar 頻譜現金（2026Q4 0.385、FY2027 9.328）列投資活動（按成本、價值中性）；"
            "現金低於最低現金 30 時以新債補足。Valor 失敗售後租回的新增借款與本金償還（債務到期表）列在融資，與其他債務同處理（10-30 更正）")
CR = {s: {} for s in "ABC"}
r = 4
CF.cell(r, 1, "收支關卡（地圖讀這兩列；FY2026–FY2030，FY2026 含 9M 實績）").font = BOLD; r += 1
hdr(CF, r, ["項目", "情境 A", "情境 B", "情境 C", "說明"]); r += 1
GATE = r; r += 7
def cput(s, key, label, unit, vals, fmt="#,##0.000", note=None, **kw):
    global r
    put(CF, r, label, unit, vals, fmt, note, **kw); CR[s][key] = r; r += 1
for s in "ABC":
    R = SR[s]; P = PR_[s]; C = CR[s]
    CF.cell(r, 1, SCN[s][2]).font = BOLD; CF.cell(r, 1).fill = SUB; r += 1
    hdr(CF, r, MODEL_HDR); r += 1
    P3 = ["1H"] + FC + ["26"]
    cput(s, "ni", "淨利", "US$B", {p: f"=PL!{COL[p]}{P['ni']}" for p in P3})
    cput(s, "da", "折舊攤銷", "US$B", {p: f"=PL!{COL[p]}{P['da']}" for p in P3})
    cput(s, "sbc", "股權報酬（非現金；含 Cursor 收購後服務期獎酬）", "US$B", {p: f"=PL!{COL[p]}{P['sbc']}" for p in P3})
    cput(s, "wc", "營運資金變動（流入＋）", "US$B", {"2H": f"=wc_k*(PL!F{P['rev']}-({QA('rev', '25')}-{QA('rev', 'm25')}))",
         **{p: f"=wc_k*(PL!{COL[p]}{P['rev']}-PL!{COL[prevfy(p)]}{P['rev']})" for p in ANN}, "26": f"=F{r}"},
         note="Q4 26＝wc_k ×(Q4 26 營收 − 2025Q4 營收〔FY2025 二手 − 9M25〕)；之後＝營收增量 × wc_k")
    cput(s, "oth", "其他（9M26 實績差額：營運資金、遞延稅、清償損失等；預測期 0）", "US$B",
         {"1H": f"=cfo_9m-E{C['ni']}-E{C['da']}-E{C['sbc']}", **{p: 0 for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "cfo", "營運現金流", "US$B", {p: f"={COL[p]}{C['ni']}+{COL[p]}{C['da']}+{COL[p]}{C['sbc']}+N({COL[p]}{C['wc']})+{COL[p]}{C['oth']}" for p in P3}, bold=True)
    cput(s, "cx_sp", "資本支出：Space", "US$B", {"1H": "=" + QA("capex_space", "1H"), **{p: f"=Seg!{COL[p]}{R['cx_sp']}" for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "cx_cn", "資本支出：Connectivity", "US$B", {"1H": "=" + QA("capex_conn", "1H"), **{p: f"=Seg!{COL[p]}{R['cx_cn']}" for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "cx_si", "資本支出：SI（現金）", "US$B", {p: f"=Thesis!{COL[p]}{T['si_cash']}" for p in P3})
    cput(s, "cx", "資本支出合計（現金）", "US$B", {p: f"={COL[p]}{C['cx_sp']}+{COL[p]}{C['cx_cn']}+{COL[p]}{C['cx_si']}" for p in P3}, bold=True)
    cput(s, "chk", "檢查：9M26 資本支出 − 10-Q 購置不動產與設備", "US$B", {"1H": f"=E{C['cx']}-capex_9m"}, "0.000")
    cput(s, "fin", "以其他融資取得 PP&E（非現金；SI）", "US$B", {p: f"=Thesis!{COL[p]}{T['si_fin']}" for p in P3})
    cput(s, "lease", "融資租賃本金", "US$B", {"1H": "=lease_1H", **{p: f"=lease_{p}" for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "slb", "債務到期償還（含 Valor 失敗售後租回本金；融資活動，不扣 FCF）", "US$B", {"1H": "=slb_1H", **{p: f"=mat_{p}" for p in FC}, "26": f"=E{r}+F{r}"},
         note="10-30 更正：Valor 失敗售後租回資產原以現金購置、已在資本支出內，付款以債務入帳 → 本金償還列融資來源（與其他債務同處理），不扣 FCF。9M26 未單獨揭露（0）；預測期＝10-Q 到期表")
    cput(s, "fcf_co", "公司口徑 FCF（營運現金 − 資本支出）", "US$B", {p: f"={COL[p]}{C['cfo']}-{COL[p]}{C['cx']}" for p in P3})
    cput(s, "fcf", "自由現金流（收支關卡口徑：− 融資租賃本金 − 以其他融資取得 PP&E）", "US$B",
         {p: f"={COL[p]}{C['fcf_co']}-{COL[p]}{C['lease']}-{COL[p]}{C['fin']}" for p in P3}, bold=True, fill=SUB)
    cput(s, "spec", "EchoStar 頻譜（投資活動；按成本）", "US$B", {"1H": 0, "2H": "=spec_2h", 27: "=spec_27", **{p: 0 for p in ANN[1:]}, "26": f"=E{r}+F{r}"})
    cput(s, "dnet", "9M26 債務淨流入／預測期已知新債", "US$B", {"1H": "=debt_1H", **{p: 0 for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "eq", "股權淨流入（9M26 IPO 等）", "US$B", {"1H": "=eq_1H", **{p: 0 for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "copen", "期初現金（現金＋有價證券）", "US$B", {"1H": "=cash_ye25", **{p: f"={COL[prevp(p)]}{r + 4}" for p in FC}})
    C["oth2"] = r; r += 1
    cput(s, "pre", "新債前現金", "US$B", {"1H": f"=E{C['copen']}+E{C['cfo']}-E{C['cx']}-E{C['lease']}+E{C['dnet']}+E{C['eq']}+E{C['oth2']}",
         **{p: f"={COL[p]}{C['copen']}+{COL[p]}{C['cfo']}-{COL[p]}{C['cx']}-{COL[p]}{C['lease']}-{COL[p]}{C['slb']}-{COL[p]}{C['spec']}" for p in FC}})
    cput(s, "plug", "新增舉債（現金低於最低現金時）", "US$B", {"1H": 0, **{p: f"=MAX(0,min_cash-{COL[p]}{C['pre']})" for p in FC}, "26": f"=E{r}+F{r}"})
    cput(s, "cend", "期末現金（現金＋有價證券）", "US$B", {p: f"={COL[p]}{C['pre']}+{COL[p]}{C['plug']}" for p in ["1H"] + FC})
    assert C["cend"] == C["copen"] + 4
    put(CF, C["oth2"], "其他（9M26 實績差額：收購現金、產品回饋金、匯率、頻譜預付等；預測期 0）", "US$B",
        {"1H": f"=cash_0-(E{C['copen']}+E{C['cfo']}-E{C['cx']}-E{C['lease']}+E{C['dnet']}+E{C['eq']})", **{p: 0 for p in FC}})
    cput(s, "chk2", "檢查：9M26 期末現金 − 2026-09-30 資產負債表", "US$B", {"1H": f"=E{C['cend']}-cash_0"}, "0.000")
    cput(s, "dopen", "期初債務本金", "US$B", {"1H": "=debt_ye25", **{p: f"={COL[prevp(p)]}{r + 1}" for p in FC}})
    cput(s, "dend", "期末債務本金（＋以其他融資取得 PP&E）", "US$B", {"1H": "=debt_0", **{p: f"={COL[p]}{C['dopen']}-{COL[p]}{C['slb']}+{COL[p]}{C['plug']}+{COL[p]}{C['fin']}" for p in FC}})
    cput(s, "nc", "淨現金（現金 − 債務）", "US$B", {p: f"={COL[p]}{C['cend']}-{COL[p]}{C['dend']}" for p in ["1H"] + FC})
    r += 1
widths(CF)
CF.freeze_panes = "C4"
GL = ["收支關卡：自由現金流（收支關卡口徑，五年合計）", "  對照：公司口徑 FCF（營運現金 − 資本支出，五年合計）", "  其中：五年新增舉債", "  期末淨現金 FY2030",
      "是否需要舉債", "檢查：收支關卡 − 首年起五年逐欄加總（FY2026＋…＋FY2030）"]
for i, lab in enumerate(GL): CF.cell(GATE + i, 1, lab)
for j, s in enumerate("ABC"):
    c = "BCD"[j]; C = CR[s]
    CF[f"{c}{GATE}"] = f"=SUM(G{C['fcf']}:K{C['fcf']})"
    CF[f"{c}{GATE + 1}"] = f"=SUM(G{C['fcf_co']}:K{C['fcf_co']})"
    CF[f"{c}{GATE + 2}"] = f"=SUM(G{C['plug']}:K{C['plug']})"
    CF[f"{c}{GATE + 3}"] = f"=K{C['nc']}"
    for k_ in range(4): CF[f"{c}{GATE + k_}"].number_format = "#,##0.0"
    CF[f"{c}{GATE + 4}"] = f'=IF({c}{GATE + 2}>0.001,"需要舉債","不需舉債（IPO 現金足以支應）")'
    CF[f"{c}{GATE + 5}"] = f"={c}{GATE}-(G{C['fcf']}+H{C['fcf']}+I{C['fcf']}+J{C['fcf']}+K{C['fcf']})"; CF[f"{c}{GATE + 5}"].number_format = "0.000"
CF.cell(GATE, 5, "審查第 2 條口徑（10-30 更正：不扣 Valor 本金）；負值＝入不敷出（由 IPO 現金與舉債支應）")
CF.cell(GATE + 1, 5, "公司口徑（不扣融資租賃本金、以融資取得 PP&E）")
for s in "ABC":
    P = PR_[s]; C = CR[s]
    put(PL, P["ii"], "利息收入（期初現金 × r_cash）", "US$B", {**{p: f"=r_cash*CF!{COL[p]}{C['copen']}*{LEN[p]}" for p in FC}, "26": f"=F{P['ii']}"})
    put(PL, P["ie"], "利息費用（期初債務 × r_debt）", "US$B", {**{p: f"=r_debt*CF!{COL[p]}{C['dopen']}*{LEN[p]}" for p in FC}, "26": f"=F{P['ie']}"})
    v = {"1H": "=sh_out"}
    for p in FC: v[p] = f"={COL[prevp(p)]}{P['she']}*(1+dil*{LEN[p]})"
    put(PL, P["she"], "期末股數（十億股）", "十億股", v, "0.000", "9M26 期末＝10-Q 封面 A＋B 股；之後 ×(1＋dil × 期間長度)；不買回")
    v = {"1H": "=sh_dil_1H", "26": f"=0.75*E{P['sh']}+0.25*F{P['sh']}"}
    for p in FC: v[p] = f"=({COL[prevp(p)]}{P['she']}+{COL[p]}{P['she']})/2"
    put(PL, P["sh"], "平均稀釋股數（十億股）", "十億股", v, "0.000", "9M26 為 IPO 前後加權（10-Q）")

# =====================================================================================================
# Val：0.0 折現率、0.1 DCF、0.15 SI 對外出租模組 DCF、0.2 分部加總、0.3 結論
V["A1"] = ("Val — 目標價（2027-10-30）＝DCF 50%＋分部加總（SOTP）50%；情境權重 A 25%／B 50%／C 25%（In 頁，黃底＝判斷題）。分部加總即本模型的 Comparable（在分部層做）。"
           "主值＝自用 AI 讀法 2（In 頁 internal_ai_reading），讀法 1 並列")
V["A2"] = ("DCF：無槓桿自由現金流（NOPAT＋折舊 − 全部投資〔現金資本支出＋以其他融資取得 PP&E〕＋營運資金；股權報酬視為成本、不加回）以 WACC 折現到評價日 2026-09-30（期中慣例），"
           "終值＝NOPAT(FY2033) ×(1 − g ÷ RONIC) ÷ (WACC − g)，以 ke 推到目標時點。SOTP：目標時點 NTM 分部調整後 EBITDA × 同業倍數；SI 對外出租用自身 DCF；Space 用 EV／營收")
VR_ = {}
SENS = json.load(open(A.sens, encoding="utf-8")) if A.sens else {}
r = 4
COLS3 = "BCD"
def vput(key, label, vals, fmt="#,##0.00", note=None, bold=False, fill=None):
    global r
    V.cell(r, 1, label)
    if bold: V.cell(r, 1).font = BOLD
    for j, v in enumerate(vals):
        if v is None: continue
        c = V.cell(r, 2 + j, v); c.number_format = fmt[j] if isinstance(fmt, (list, tuple)) else fmt
    if note: V.cell(r, 5, note)
    if fill:
        for cc in range(1, 6): V.cell(r, cc).fill = fill
    if key: VR_[key] = r
    r += 1
S3 = lambda f: [f(s) for s in "ABC"]
X3 = lambda s: COLS3["ABC".index(s)]
V.cell(r, 1, "0.3 結論").font = BOLD; r += 1
hdr(V, r, ["項目", "數值", "對照", "對照", "說明"]); r += 1
CONC = r; r += 27
V.cell(r, 1, "0.0 折現率（CAPM；負債含融資租賃）").font = BOLD; r += 1
hdr(V, r, ["項目", "數值", "", "", "說明"]); r += 1
vput("emkt", "股權市值 E＝評價日股價 × 流通股", ["=px_val*sh_out"], "#,##0.0")
vput("dmkt", "負債 D（債務本金＋融資租賃）", ["=debt_0+fin_lease"], "#,##0.0")
vput("ke", "股東權益成本 ke＝rf＋β × ERP", ["=rf+beta*erp"], "0.00%", note='="β（判斷）"&TEXT(beta,"0.00")')
dn("ke", f"Val!$B${VR_['ke']}")
vput("kdat", "稅後債務成本", ["=kd*(1-tax_w)"], "0.00%")
vput("wacc", "WACC（＋wacc_adj，Sens 用）", [f"=(B{VR_['emkt']}*ke+B{VR_['dmkt']}*B{VR_['kdat']})/(B{VR_['emkt']}+B{VR_['dmkt']})+wacc_adj"], "0.00%", bold=True)
dn("wacc", f"Val!$B${VR_['wacc']}")
vput("ronic", "永續期新投資報酬 RONIC＝WACC＋ronic_sp", ["=wacc+ronic_sp"], "0.00%")
dn("ronic", f"Val!$B${VR_['ronic']}")
r += 1
V.cell(r, 1, "0.1 DCF（無槓桿自由現金流；三情境各一區塊；欄 F＝Q4 26、H–M＝FY2027–FY2032）＋0.15 SI 對外出租模組 DCF（分部加總的對外出租腿）").font = BOLD; r += 1
DC = {s: {} for s in "ABC"}
VC = {"2H": "F", 27: "H", 28: "I", 29: "J", 30: "K", 31: "L", 32: "M"}
for s in "ABC":
    P = PR_[s]; C = CR[s]; D = DC[s]; R = SR[s]
    V.cell(r, 1, SCN[s][2]).font = BOLD; V.cell(r, 1).fill = SUB; r += 1
    hdr(V, r, ["項目", "", "", "", "說明", YH["2H"], "", YH[27], YH[28], YH[29], YH[30], YH[31], YH[32]]); r += 1
    def drow(key, label, f, fmt="#,##0.00", note=None):
        global r
        V.cell(r, 1, label)
        for p in FC:
            c = V[f"{VC[p]}{r}"]; c.value = f(p); c.number_format = fmt
        if note: V.cell(r, 5, note)
        D[key] = r; r += 1
    drow("ebit", "營業利益（財報口徑 EBIT）", lambda p: f"=PL!{COL[p]}{P['oi']}")
    drow("nopat", "NOPAT＝(EBIT＋sbc_add × 股權報酬) ×(1 − 稅率)", lambda p: f"=({VC[p]}{D['ebit']}+sbc_add*PL!{COL[p]}{P['sbc']})*(1-tax)")
    drow("da", "＋折舊攤銷", lambda p: f"=PL!{COL[p]}{P['da']}")
    drow("cx", "− 投資（現金資本支出＋以其他融資取得 PP&E）", lambda p: f"=CF!{COL[p]}{C['cx']}+CF!{COL[p]}{C['fin']}")
    drow("wc", "＋營運資金變動", lambda p: f"=CF!{COL[p]}{C['wc']}")
    drow("ufcf", "無槓桿自由現金流（財報口徑）", lambda p: f"={VC[p]}{D['nopat']}+{VC[p]}{D['da']}-{VC[p]}{D['cx']}+{VC[p]}{D['wc']}")
    drow("shd", "自用 AI 影子效益（稅後；讀法 2 加入）", lambda p: f"=PL!{COL[p]}{P['shadow']}*(1-tax)")
    drow("tvn", "終值用 NOPAT（終值期稅率 tax_tv）", lambda p: f"=({VC[p]}{D['ebit']}+sbc_add*PL!{COL[p]}{P['sbc']})*(1-tax_tv)")
    drow("tvf", "終值用無槓桿自由現金流（tax_tv；tv_alt＝1 時用）", lambda p: f"={VC[p]}{D['tvn']}+{VC[p]}{D['da']}-{VC[p]}{D['cx']}+{VC[p]}{D['wc']}")
    drow("shdt", "終值用影子效益（tax_tv）", lambda p: f"=PL!{COL[p]}{P['shadow']}*(1-tax_tv)")
    drow("t", "折現年數（評價日 2026-09-30 起，期中）", lambda p: "=t_q4" if p == "2H" else f"={PY[p]}+mid-t_val", "0.00")
    drow("df", "折現因子 1 ÷ (1＋WACC)^年數", lambda p: f"=1/(1+wacc)^{VC[p]}{D['t']}", "0.0000")
    drow("pv", "現值（財報口徑）", lambda p: f"={VC[p]}{D['ufcf']}*{VC[p]}{D['df']}")
    drow("pvs", "現值（影子效益）", lambda p: f"={VC[p]}{D['shd']}*{VC[p]}{D['df']}")
    drow("r_e", "0.15 SI 對外出租 EBITDA", lambda p: f"=PL!{COL[p]}{P['e_rent']}")
    drow("r_d", "0.15 對外新增 AI 折舊", lambda p: f"=Thesis!{COL[p]}{T['da_ext']}")
    drow("r_np", "0.15 對外出租 NOPAT", lambda p: f"=({VC[p]}{D['r_e']}-{VC[p]}{D['r_d']})*(1-tax)")
    drow("r_npt", "0.15 對外出租終值用 NOPAT（tax_tv）", lambda p: f"=({VC[p]}{D['r_e']}-{VC[p]}{D['r_d']})*(1-tax_tv)")
    drow("r_cx", "0.15 對外出租投資（SI 投資 × ext_sh）", lambda p: f"=Thesis!{COL[p]}{T['si_inv']}*ext_sh")
    drow("r_f", "0.15 對外出租自由現金流", lambda p: f"={VC[p]}{D['r_np']}+{VC[p]}{D['r_d']}-{VC[p]}{D['r_cx']}")
    drow("r_pv", "0.15 現值", lambda p: f"={VC[p]}{D['r_f']}*{VC[p]}{D['df']}")
    r += 1
V.cell(r, 1, "DCF 彙總").font = BOLD; r += 1
hdr(V, r, ["項目", "情境 A", "情境 B", "情境 C", "說明"]); r += 1
TV_T = lambda s: f"($M${DC[s]['t']}+mid)"
vput("pvx", "明確期現值合計（財報口徑）", S3(lambda s: f"=SUM(F{DC[s]['pv']}:M{DC[s]['pv']})"))
vput("tvv", "終值（FY2032 年底，財報口徑）", S3(lambda s: f"=IF(tv_alt=1,M{DC[s]['tvf']}*(1+g_tv)/(wacc-g_tv),M{DC[s]['tvn']}*(1+g_tv)*(1-g_tv/ronic)/(wacc-g_tv))"),
     note="價值驅動：NOPAT(FY2033，終值期稅率 21%) ×(1 − g ÷ RONIC) ÷ (WACC − g)")
vput("pvtv", "終值現值", S3(lambda s: f"={X3(s)}{VR_['tvv']}/(1+wacc)^{TV_T(s)}"))
vput("ev1", "企業價值 EV（讀法 1：自用 AI 只計成本）", S3(lambda s: f"={X3(s)}{VR_['pvx']}+{X3(s)}{VR_['pvtv']}"))
vput("pvsx", "＋自用 AI 影子效益現值：明確期", S3(lambda s: f"=SUM(F{DC[s]['pvs']}:M{DC[s]['pvs']})"))
vput("pvstv", "＋自用 AI 影子效益現值：終值", S3(lambda s: f"=IF(tv_alt=1,M{DC[s]['shdt']}*(1+g_tv)/(wacc-g_tv),M{DC[s]['shdt']}*(1+g_tv)*(1-g_tv/ronic)/(wacc-g_tv))/(1+wacc)^{TV_T(s)}"))
vput("ev2", "企業價值 EV（讀法 2：自用 AI 價值中性）", S3(lambda s: f"={X3(s)}{VR_['ev1']}+{X3(s)}{VR_['pvsx']}+{X3(s)}{VR_['pvstv']}"))
vput("ev", "企業價值 EV（主值，依 internal_ai_reading）", S3(lambda s: f"=IF(internal_ai_reading=2,{X3(s)}{VR_['ev2']},{X3(s)}{VR_['ev1']})"), bold=True)
vput("tvsh", "  其中：終值占 EV（主值）", S3(lambda s: f"=({X3(s)}{VR_['pvtv']}+(internal_ai_reading=2)*{X3(s)}{VR_['pvstv']})/{X3(s)}{VR_['ev']}"), "0.0%")
vput("nc0", "＋淨現金（評價日：現金＋有價證券 − 債務本金 − 融資租賃）", ["=cash_0-debt_0-fin_lease"] * 3)
vput("hold", "＋數位資產（比特幣帳面）", ["=btc"] * 3)
vput("dl", "− 類債（訴訟準備）", ["=legal"] * 3, note="EchoStar 頻譜：未來付款與取得的頻譜按成本相抵（價值中性），DCF 兩邊都不計")
vput("shv", "股數（評價日 A＋B；dil_x＝1 時加潛在稀釋）", ["=sh_out+dil_x*(dil_perf+dil_opt+dil_rsu+dil_echo)"] * 3, "0.000")
BR = lambda ev: (lambda s: "=MAX(0,({0}{1}+{0}{2}+{0}{3}-{0}{4})/{0}{5}*(1+ke)^t_tgt)".format(X3(s), VR_[ev], VR_["nc0"], VR_["hold"], VR_["dl"], VR_["shv"]))
vput("dcf1", "DCF 每股價值推到目標時點（讀法 1）", S3(BR("ev1")), note="(EV＋淨現金＋數位資產 − 類債) ÷ 股數 ×(1＋ke)^t_tgt；不配息。情境 C 權益為負，被 MAX(0) 截為 0；讀法 1 情境 B 的 EV 低於 Connectivity 分部加總 EV → DCF 隱含 SI 整體價值為負")
vput("dcf2", "DCF 每股價值推到目標時點（讀法 2）", S3(BR("ev2")))
vput("dcf12", "DCF 情境價值推到 12 個月後（主值）", S3(lambda s: f"=IF(internal_ai_reading=2,{X3(s)}{VR_['dcf2']},{X3(s)}{VR_['dcf1']})"), bold=True, fill=KEY)
vput("r_pv", "0.15 SI 對外出租明確期現值（FY2028–FY2032）", S3(lambda s: f"=SUM(I{DC[s]['r_pv']}:M{DC[s]['r_pv']})"), note="Q4 26–FY2027 已反映在 FY2027 年底淨現金")
vput("r_tv", "0.15 SI 對外出租終值現值", S3(lambda s: f"=M{DC[s]['r_npt']}*(1+g_tv)*(1-g_tv/ronic)/(wacc-g_tv)/(1+wacc)^{TV_T(s)}"))
vput("r_v", "0.15 SI 對外出租價值（FY2027 年底，與分部加總淨現金同時點）", S3(lambda s: f"=({X3(s)}{VR_['r_pv']}+{X3(s)}{VR_['r_tv']})*(1+wacc)^($H${DC[s]['t']}+mid)"), bold=True,
     note="評價日對外機隊的資本已投入（不再扣），其 EBITDA 計入；之後新增與汰換的對外投資全數扣除")
r += 1
V.cell(r, 1, "0.2 分部加總（SOTP；目標時點 NTM 分部調整後 EBITDA × 同業 EV/EBITDA；＝本模型的 Comparable）").font = BOLD; r += 1
hdr(V, r, ["同業組／倍數", "倍數", "家數", "", "說明"]); r += 1
ck = [f"pm_conn_{p_['ticker'].lower()}" for p_ in PCN["peers"] if p_.get("evEbitdaFy2026") is not None]
sk = [f"pm_si_{p_['ticker'].lower()}" for p_ in PSI["peers"]]
vput("mult_conn", "Connectivity：衛星與電信同業 EV/EBITDA FY2026 中位數 × mult_cn_adj × mult_cn_pm", ["=MEDIAN(" + ",".join(ck) + ")*mult_cn_adj*mult_cn_pm", len(ck)], ["0.00", "0"], note="data/peers_connectivity_20261030.json（2026-10-08）：" + "、".join(k.split("_")[-1].upper() for k in ck))
dn("mult_conn", f"Val!$B${VR_['mult_conn']}")
vput("mult_si", "自用 AI 營運成本加回倍數：AI 基礎設施／雲同業中位數 × mult_si_adj", ["=MEDIAN(" + ",".join(sk) + ")*mult_si_adj", len(sk)], ["0.00", "0"],
     note="data/peers_si_20261030.json：CRWV、NBIS FY2027、ORCL NTM（口徑不一、樣本小；README 註明）")
dn("mult_si", f"Val!$B${VR_['mult_si']}")
r += 1
hdr(V, r, ["分部（NTM 分部調整後 EBITDA；Space 為營收）", "情境 A", "情境 B", "情境 C", "倍數與說明", "EV：A", "EV：B", "EV：C"]); r += 1
NT = lambda s, k: f"(w_ntm27*PL!H{PR_[s][k]}+(1-w_ntm27)*PL!I{PR_[s][k]})"
NTT = lambda k: f"(w_ntm27*Thesis!H{T[k]}+(1-w_ntm27)*Thesis!I{T[k]})"
LINES = [("cn", "Connectivity（分部調整後 EBITDA）", lambda s: f"={NT(s, 'e_cn')}", "mult_conn", "衛星與電信"),
         ("so", "SI 其他：Grok／Cursor／X 與廣告（NTM 營收 × EV／營收）", lambda s: f"={NT(s, 'nr')}+{NT(s, 'ads')}", "evs_si", "EV／營收（判斷）"),
         ("sp", "Space（NTM 營收 × EV／營收）", lambda s: f"={NT(s, 'space')}", "evs_space", "EV／營收（判斷）")]
SOTP0 = r
for key, lab, f, mult, mlab in LINES:
    V.cell(r, 1, lab)
    for j, s in enumerate("ABC"):
        V.cell(r, 2 + j, f(s)).number_format = "#,##0.00"
        V.cell(r, 6 + j, f"={COLS3[j]}{r}*{mult}").number_format = "#,##0.0"
    V.cell(r, 5, f"=\"{mlab} \"&TEXT({mult},\"0.00\")&\"×\"")
    VR_["s_" + key] = r; r += 1
V.cell(r, 1, "SI 對外雲端服務（EV＝0.15 自身 DCF；EBITDA 只供參考）")
for j, s in enumerate("ABC"):
    V.cell(r, 2 + j, f"={NT(s, 'e_rent')}").number_format = "#,##0.00"
    V.cell(r, 6 + j, f"={COLS3[j]}{VR_['r_v']}").number_format = "#,##0.0"
V.cell(r, 5, "對外出租模組 DCF"); VR_["s_rent"] = r; r += 1
SOTP1 = r - 1
vput("sev1", "分部加總 EV（讀法 1）", [f"=SUM(F{SOTP0}:F{SOTP1})", f"=SUM(G{SOTP0}:G{SOTP1})", f"=SUM(H{SOTP0}:H{SOTP1})"], "#,##0.0")
V.cell(r, 1, "＋自用 AI 增量營運成本加回（讀法 2；價值中性）")
for j, s in enumerate("ABC"):
    V.cell(r, 2 + j, f"={NTT('int_opex')}").number_format = "#,##0.00"
    V.cell(r, 6 + j, f"={COLS3[j]}{r}*mult_si").number_format = "#,##0.0"
V.cell(r, 5, "同 Alphabet 修正二（自用 AI 的資本成本不在 EBITDA，故不加影子效益）"); VR_["s_sh"] = r; r += 1
V.cell(r, 1, "＋自用 AI 新增資產帳面淨值（FY2027 年底；讀法 2 價值中性＝按帳面）")
_nbv = f"=(1-ext_sh)*(cip_0*cip_si+Thesis!F{T['si_inv']}+Thesis!H{T['si_inv']})-(Thesis!F{T['da_int']}+Thesis!H{T['da_int']})"
for j, s in enumerate("ABC"):
    V.cell(r, 6 + j, _nbv).number_format = "#,##0.0"
V.cell(r, 5, "FY2027 年底淨現金已扣 Q4 26–FY2027 自用 AI 投資與在建工程自用部分，讀法 2 下這些資產按帳面計"); VR_["s_nbv"] = r; r += 1
vput("sev2", "分部加總 EV（讀法 2）", [f"=B{VR_['sev1']}+F{VR_['s_sh']}+F{VR_['s_nbv']}", f"=C{VR_['sev1']}+G{VR_['s_sh']}+G{VR_['s_nbv']}", f"=D{VR_['sev1']}+H{VR_['s_sh']}+H{VR_['s_nbv']}"], "#,##0.0")
vput("sev", "分部加總 EV（主值）", S3(lambda s: f"=IF(internal_ai_reading=2,{X3(s)}{VR_['sev2']},{X3(s)}{VR_['sev1']})"), "#,##0.0", bold=True)
vput("snc", "＋淨現金（FY2027 年底，CF 頁；減融資租賃）", S3(lambda s: f"=CF!H{CR[s]['nc']}-fin_lease"))
vput("shold", "＋數位資產＋EchoStar 頻譜（按 Q4 26–FY2027 已付成本）", ["=btc+spec_2h+spec_27"] * 3, note="頻譜付款已自淨現金扣除，資產按成本加回（價值中性）")
vput("sdl", "− 類債", ["=legal"] * 3)
vput("ssh", "股數（FY2027 年底；dil_x＝1 時加潛在稀釋）", S3(lambda s: f"=PL!H{PR_[s]['she']}+dil_x*(dil_perf+dil_opt+dil_rsu+dil_echo)"), "0.000")
SB2 = lambda ev: (lambda s: "=MAX(0,({0}{1}+{0}{2}+{0}{3}-{0}{4})/{0}{5})".format(X3(s), VR_[ev], VR_["snc"], VR_["shold"], VR_["sdl"], VR_["ssh"]))
vput("sotp1", "分部加總每股價值（讀法 1）", S3(SB2("sev1")))
vput("sotp2", "分部加總每股價值（讀法 2）", S3(SB2("sev2")))
vput("sotp", "分部加總每股價值（主值；目標時點）", S3(lambda s: f"=IF(internal_ai_reading=2,{X3(s)}{VR_['sotp2']},{X3(s)}{VR_['sotp1']})"), bold=True, fill=KEY)
vput("sebt", "  分部 NTM 調整後 EBITDA 合計（對照）", S3(lambda s: f"={NT(s, 'ebitda')}"), "#,##0.0")
vput("simp", "  隱含 EV ÷ NTM 調整後 EBITDA（主值）", S3(lambda s: "={0}{1}/{0}{2}".format(X3(s), VR_["sev"], VR_["sebt"])), "0.00")
vput("mkt_imp", "  對照：現價隱含 EV ÷ NTM 調整後 EBITDA（情境 B）", [f"=(px*sh_out-cash_0+debt_0+fin_lease)/C{r - 1 - 0}"], "0.00")
V[f"B{VR_['mkt_imp']}"] = f"=(px*sh_out-cash_0+debt_0+fin_lease)/C{VR_['sebt']}"
r_end = r; r = CONC
W3 = lambda row: f"tp_wA*B{row}+tp_wB*C{row}+tp_wC*D{row}"
vput("tpdcf", "DCF 目標價（情境加權）", [f"={W3(VR_['dcf12'])}"], "#,##0.00", fill=KEY)
vput("tpcomp", "Comparable 目標價（市場主流同業法）", [f"={W3(VR_['sotp'])}"], "#,##0.00", note="＝分部加總（SOTP）情境加權；Connectivity、SI 其他用同業倍數，SI 對外出租用自身 DCF，Space 用 EV／營收")
vput("tp", "目標價（DCF 與 Comparable 加權）", ["=w_dcf*B{0}+(1-w_dcf)*B{1}".format(r - 2, r - 1)], "#,##0.00", bold=True, fill=KEY,
     note='="DCF 權重 "&TEXT(w_dcf,"0%")&"、分部加總 "&TEXT(1-w_dcf,"0%")&"；自用 AI 讀法 "&internal_ai_reading')
vput("px", "現價（2026-10-29 收盤）", ["=px"], "#,##0.00", note="MarketBeat、Equibles 兩來源一致（In 頁 px）")
vput("cons", "市場共識目標價", ["=cons_tp"], "#,##0.00", note="MarketBeat 46 位分析師平均（含 10-21～10-26 財報後調整）；只作對照")
vput("up", "上檔空間：加權目標價", [f"=B{VR_['tp']}/px-1"], "0.0%", bold=True, fill=KEY)
vput("upc", "上檔空間：Comparable 目標價", [f"=B{VR_['tpcomp']}/px-1"], "0.0%")
def tier(u): return f'IF({u}>=0.5,"大於 50%",IF({u}>=0.25,"25%–50%",IF({u}>=0,"0%–25%",IF({u}>-0.25,"0%～−25%",IF({u}>-0.5,"−25%～−50%","小於 −50%")))))'
def recf(u): return f'IF({u}>=0.5,"強力買進",IF({u}>=0.25,"買進",IF({u}>-0.25,"中立",IF({u}>-0.5,"放空","強力放空"))))'
U = f"B{VR_['up']}"
vput("tier", "上檔空間級距", ["=" + tier(U)], "@", note="以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
vput("rec", "投資建議", ["=" + recf(U)], "@", bold=True, fill=KEY, note="五級固定規則：≥ +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、≤ −50% 強力放空")
vput("band", "價格帶（DCF 不如預期～順利）", [f"=D{VR_['dcf12']}", f"=B{VR_['dcf12']}"], "#,##0.00", note="B＝C 情境 DCF、C＝A 情境 DCF（皆推到 12 個月後）")
vput("pess", "悲觀情境價（C 情境 DCF）", [f"=D{VR_['dcf12']}"], "#,##0.00", note="地圖「悲觀情境空間」＝本列 ÷ 現價 − 1")
vput("scen", "情境目標價（A／B／C；DCF 與分部加總加權）", S3(lambda s: "=w_dcf*{0}{1}+(1-w_dcf)*{0}{2}".format(X3(s), VR_["dcf12"], VR_["sotp"])), "#,##0.00")
vput("scenup", "情境上檔空間（A／B／C）", S3(lambda s: "={0}{1}/px-1".format(X3(s), VR_["scen"])), "0.0%")
vput("tvsh2", "DCF 終值占 EV（A／B／C）", S3(lambda s: "={0}{1}".format(X3(s), VR_["tvsh"])), "0.0%")
TPR = lambda d, sp: f"=w_dcf*({W3(VR_[d])})+(1-w_dcf)*({W3(VR_[sp])})"
vput("tp_r2", "讀法 2（自用 AI 價值中性）：目標價／上檔空間／投資建議", [TPR("dcf2", "sotp2"), f"=B{r}/px-1", "=" + recf(f"C{r}")], ["#,##0.00", "0.0%", "@"], note="活公式（主值）")
vput("tp_r1", "讀法 1（自用 AI 只計成本）：目標價／上檔空間／投資建議", [TPR("dcf1", "sotp1"), f"=B{r}/px-1", "=" + recf(f"C{r}")], ["#,##0.00", "0.0%", "@"], note="活公式（下緣敏感度）")
vput("tp_dcf12", "讀法 1／2 的 DCF 目標價（情境加權）", [f"={W3(VR_['dcf1'])}", f"={W3(VR_['dcf2'])}"], "#,##0.00")
vput("tp_sotp12", "讀法 1／2 的分部加總目標價（情境加權）", [f"={W3(VR_['sotp1'])}", f"={W3(VR_['sotp2'])}"], "#,##0.00")
vput("s2", "S2 核心問題：FY2028 SI 每 GW 年營收 ÷ Tokenomics 每 GW 年持有成本（A／B／C）", S3(lambda s: f"=Seg!I{SR[s]['s2']}"), "0.00",
     note='="假設路徑內插、非模型產出；2026Q3 校準 k0＝"&TEXT(k0,"0.00")&"；FY2028 情境 B 每 GW 年營收 "&TEXT(Seg!I' + str(SR["B"]["rev_gw"]) + ',"0.0")&"、持有成本 "&TEXT(Thesis!I' + str(T["he_fl"]) + ',"0.0")&"（US$B/GW-年）"')
vput("gate", "收支關卡：FY2026–FY2030 自由現金流（收支關卡口徑；A／B／C）", S3(lambda s: f"=CF!{X3(s)}{GATE}"), "#,##0.0", note="CF 頁；公司口徑對照見 CF 頁")
vput("meth", "目標價方法", ["DCF（無槓桿 FCF、WACC、RONIC＝WACC）50%＋分部加總（Connectivity、SI 其他同業倍數；SI 對外出租自身 DCF；Space EV／營收）50%；情境 25／50／25；自用 AI 讀法 2"], "@")
vput("wsum", "檢查：情境權重合計（應為 1）", ["=ROUND(tp_wA+tp_wB+tp_wC,6)"], "0.000")
vput("chk_rd", "檢查：主值目標價 − 所選讀法的目標價（應為 0）", [f"=B{VR_['tp']}-IF(internal_ai_reading=2,B{VR_['tp_r2']},B{VR_['tp_r1']})"], "0.000")
assert r <= CONC + 27, r
r = r_end
V.column_dimensions["A"].width = 66
for c in "BCD": V.column_dimensions[c].width = 13
V.column_dimensions["E"].width = 60
for c in "FGHIJKLM": V.column_dimensions[c].width = 12

# =====================================================================================================
# Recon：口徑對照
RC["A1"] = ("Recon — 口徑對照：合併層對市場共識（Q3 後年度共識未找到，只有 2026-06-22 過期的 Simply Wall St）、逐分部對公司指引與營運數字、S2 核心問題比值的兩種讀法、"
            "每 GW 投資對照。只作對照，不作校準")
r = 3
RC.cell(r, 1, "1. 合併層：本模型 vs 過期共識（Simply Wall St，2026-06-22，IPO 後首週）").font = BOLD; r += 1
hdr(RC, r, ["項目", "年度", "情境 A", "情境 B", "情境 C", "過期共識", "B ÷ 共識 − 1", "說明"]); r += 1
RCR = {}
for k, lab, src in (("revenue", "營收", "rev"), ("earnings", "淨利", "ni"), ("fcf", "公司口徑 FCF", "fcf_co")):
    for y in (26, 27, 28):
        RC.cell(r, 1, lab); RC.cell(r, 2, f"FY20{y}")
        for j, s in enumerate("ABC"):
            col = COL["26" if y == 26 else y]
            ref = f"PL!{col}{PR_[s][src]}" if src in ("rev", "ni") else f"CF!{col}{CR[s][src]}"
            RC.cell(r, 3 + j, f"={ref}").number_format = "#,##0.00"
        RC.cell(r, 6, f"=c_{k}_{y}").number_format = "#,##0.00"
        RC.cell(r, 7, f"=D{r}/F{r}-1").number_format = "0.0%"
        if y == 26 and k == "revenue": RC.cell(r, 8, "過期：Q3 實績 14.575 遠高於當時預期（Q3 共識 12.59，Shacknews）；Q3 後年度共識未找到")
        RCR[f"{k}{y}"] = r; r += 1
RC.cell(r, 1, "市場共識目標價（MarketBeat，46 位）"); RC.cell(r, 3, "=cons_tp").number_format = "#,##0.00"
RC.cell(r, 4, f"=Val!B{VR_['tp']}").number_format = "#,##0.00"; RC.cell(r, 5, f"=D{r}/C{r}-1").number_format = "0.0%"
RC.cell(r, 8, "C＝共識、D＝本模型加權目標價、E＝差距；共識區間 75–800"); RCR["tp"] = r; r += 2
RC.cell(r, 1, "2. 逐分部：本模型（情境 B）vs 公司指引與營運數字").font = BOLD; r += 1
hdr(RC, r, ["分部／項目", "FY2026E", "FY2027E", "FY2028E", "共識", "公司指引或原文", "說明"]); r += 1
PB = PR_["B"]
for key, lab, note in (("space", "Space", "無指引"), ("conn", "Connectivity", "無指引"), ("si", "SI (Super Intelligence; formerly AI)", "無指引"),
                       ("rent", "  SI 對外雲端服務（命題子列）", "未單獨揭露；Customer B 2026Q3 5.028（衍生）"), ("nr", "  Grok／Cursor／X 訂閱、資料授權與 API（命題子列）", "未單獨揭露"),
                       ("ads", "  Advertising", "")):
    RC.cell(r, 1, lab)
    for j, p in enumerate(("26", 27, 28)): RC.cell(r, 2 + j, f"=PL!{COL[p]}{PB[key]}").number_format = "#,##0.00"
    RC.cell(r, 5, "未找到"); RC.cell(r, 6, note)
    RC.cell(r, 7, "分部共識需付費資料（未找到）" if key == "space" else ""); r += 1
RC.cell(r, 1, "SI 名目算力 2026 年底（GW）：本模型 vs 公司 Q2 法說「逾 2 GW」（二手轉述）"); RC.cell(r, 2, f"=Thesis!F{T['end']}").number_format = "0.00"
RC.cell(r, 6, "=gw_ye26_co").number_format = "0.00"; RC.cell(r, 7, "審查第 3 條：Interested-party，只作對照"); RCR["gw26"] = r; r += 1
RC.cell(r, 1, "Backlog 62.959（67% 一年內）vs 本模型 FY2027 營收（B）"); RC.cell(r, 2, "=backlog*0.67").number_format = "#,##0.0"
RC.cell(r, 3, f"=PL!H{PB['rev']}").number_format = "#,##0.0"; RC.cell(r, 7, "Backlog 未分分部；B 欄＝一年內可認列約數"); r += 1
RC.cell(r, 1, "2026Q2 新簽雲端合約 14.1（三年固定月費）年化 vs 本模型 FY2027 對外雲端服務（B）"); RC.cell(r, 2, "=cloud_q2/3").number_format = "#,##0.0"
RC.cell(r, 3, f"=PL!H{PB['rent']}").number_format = "#,##0.0"; RC.cell(r, 7, "合約年化只是新簽部分；Q3 另有新簽、未給金額"); r += 1
RC.cell(r, 1, "不可取消承諾 FY2027（含 SI 基礎設施、第三方雲、頻譜）vs 本模型 FY2027 SI 投資（B）"); RC.cell(r, 2, "=commit_27").number_format = "#,##0.0"
RC.cell(r, 3, f"=Thesis!H{T['si_inv']}").number_format = "#,##0.0"; RC.cell(r, 7, "承諾未拆分；第三方雲容量不在名目算力內（S6）→ 可能低估 SI 算力成本"); r += 1
RC.cell(r, 1, "Cursor 收購對價（GAAP，2026-08-14）vs 本模型 SI 其他 EV（情境 B；EV／NTM 營收法）"); RC.cell(r, 2, fb("cursor.consideration")).number_format = "#,##0.0"
RC.cell(r, 3, f"=Val!G{VR_['s_so']}").number_format = "#,##0.0"; RC.cell(r, 7, "Cursor 對價 271M 股 × 141.29＋替代獎酬 4.181（事實總帳 cursor.consideration）；SI 其他另含 Grok、X 訂閱與廣告"); r += 2
RC.cell(r, 1, "3. S2 核心問題比值的讀法（審查第 3、4 條；S6）").font = BOLD; r += 1
hdr(RC, r, ["項目", "情境 A", "情境 B", "情境 C", "說明"]); r += 1
RC.cell(r, 1, "FY2028 對外每 GW 年營收（IT 口徑 GW）")
for j, s in enumerate("ABC"): RC.cell(r, 2 + j, f"=Seg!I{SR[s]['rev_gw']}").number_format = "0.00"
r += 1
RC.cell(r, 1, "FY2028 在役機隊每 GW 年持有成本（Tokenomics IF_HoldEcon；IT 口徑）")
for j, s in enumerate("ABC"): RC.cell(r, 2 + j, f"=Thesis!I{T['he_fl']}").number_format = "0.00"
r += 1
RC.cell(r, 1, "比值（讀法 A：名目算力＝IT 負載，主值）")
for j, s in enumerate("ABC"): RC.cell(r, 2 + j, f"=B{r - 2}/B{r - 1}".replace("B", "BCD"[j])).number_format = "0.00"
RC.cell(r, 5, "與 Tokenomics Interface 第 1 列 IT 關鍵電力同口徑，不乘 PUE（審查第 3 條）"); RCR["s2a"] = r; r += 1
RC.cell(r, 1, "比值（讀法 B：若名目算力實為設施功率，IT ＝ GW ÷ IF_FacilityGW 1.2）")
for j, s in enumerate("ABC"): RC.cell(r, 2 + j, f"={'BCD'[j]}{r - 1}*{tk('IF_FacilityGW', 'GB300')}").number_format = "0.00"
RC.cell(r, 5, "10-Q 定義明確排除冷卻與設施負載 → 讀法 B 不採，只列對照（S6）"); r += 1
RC.cell(r, 1, "2026Q3 校準 k0（Q3 年化 SI S&I ÷ 平均在役 GW ÷ 評價日機隊持有成本）"); RC.cell(r, 2, "=k0").number_format = "0.00"
RC.cell(r, 5, "含 Grok／Cursor／X 收入的 GW 平均；對外出租與 GW 同比例時與對外每 GW 相同"); r += 2
RC.cell(r, 1, "4. 每 GW 投資對照").font = BOLD; r += 1
RC.cell(r, 1, "Tokenomics 每 GW 投資（GB300 錨 IT＋自建 × 機房）"); RC.cell(r, 2, "=cxgw26").number_format = "0.00"; r += 1
RC.cell(r, 1, "PP&E（伺服器與網通＋資料中心基礎設施）÷ 名目算力（2026-09-30）"); RC.cell(r, 2, "=ppe_gw").number_format = "0.00"
RC.cell(r, 5, "低於 Tokenomics：含 Hopper 舊機隊、機房部分租用；第三方雲與失敗售後租回資產另列"); r += 1
RC.cell(r, 1, "近四季 SI 投資（2025Q4–2026Q3；FY2025 二手 − 9M25＋9M26＋以融資取得）÷ 名目算力增量（0.6 → 1.7）")
RC.cell(r, 2, f"=({QA('capex_si', '25')}-{QA('capex_si', 'm25')}+{QA('capex_si', '1H')}+ppefin_9m)/(gw_0-Q!C{QR['gw']})").number_format = "0.00"
RC.cell(r, 5, "含在建工程（尚未在役）→ 高於每 GW 投資屬合理"); r += 1
RC.column_dimensions["A"].width = 76
for c in "BCDEFG": RC.column_dimensions[c].width = 13
RC.column_dimensions["E"].width = 40; RC.column_dimensions["H"].width = 70

# =====================================================================================================
# Sens：一次改一個輸入（整本重算後的加權目標價；make_v01.sh 產生）
SENS_CASES = [
    ("internal_ai_reading", "自用 AI 讀法 1（只計成本）", "lo", None, "讀法 2 → 讀法 1"),
    ("ext_sh", "對外出租占比", "lo", "hi", "72% → 60%／85%（審查第 4 條區間）"),
    ("pace", "SI 名目算力新增速度", "lo", "hi", "基準 × 2/3／× 1.5"),
    ("selfb", "自建機房比例", "lo", "hi", "90% → 75%／100%"),
    ("k30_B", "情境 B 的 FY2030 k", "lo", "hi", "1.25 → 1.0／1.5（只影響情境 B）"),
    ("m_so_adj", "SI 其他 EBITDA 率", "lo", "hi", "校準值 ±10pt"),
    ("g27_nr", "SI 其他收入 FY2027 年增", "lo", "hi", "50% → 25%／80%"),
    ("gl_nr", "SI 其他收入長期成長", "lo", "hi", "15% → 8%／25%"),
    ("cip_si", "在建工程屬 SI 比例（折舊）", "lo", "hi", "85% → 0／100%"),
    ("life_it", "AI IT 折舊年限", "lo", "hi", "6 → 5／7 年"),
    ("nadd_B", "情境 B 的 Starlink FY2027 淨增", "lo", "hi", "5.5M → 4.0M／7.0M（只影響情境 B）"),
    ("g_arpu", "ARPU 年變動", "lo", "hi", "−2% → −5%／0%"),
    ("m_cn", "Connectivity 調整後 EBITDA 率", "lo", "hi", "64% → 58%／70%"),
    ("cx_cn", "Connectivity 資本支出占營收", "lo", "hi", "33% → 25%／40%"),
    ("gl_eg", "企業與政府長期成長", "lo", "hi", "10% → 5%／15%"),
    ("m_spl", "Space FY2030 調整後 EBITDA 率", "lo", "hi", "20% → 5%／30%"),
    ("cx_spl", "Space FY2030 資本支出占營收", "lo", "hi", "30% → 20%／50%"),
    ("nl_27", "FY2027 客戶發射次數", "lo", "hi", "48 → 40／60"),
    ("tax", "有效稅率", "lo", "hi", "12% → 5%／21%"),
    ("wacc_adj", "WACC", "lo", "hi", "−1pt／＋1pt"),
    ("beta", "β", "lo", "hi", "1.3 → 1.1／1.6"),
    ("erp", "股權風險溢酬", "lo", "hi", "4%／6%"),
    ("g_tv", "永續成長 g", "lo", "hi", "2%／4%"),
    ("ronic_sp", "RONIC＝WACC＋5pt", None, "hi", "RONIC＝WACC → WACC＋5pt"),
    ("mult_cn_adj", "Connectivity 倍數：同業區間下緣／上緣", "lo", "hi", "7.13× → 4.71×（CMCSA）／16.66×（SATS）"),
    ("mult_cn_pm", "Connectivity 倍數 ±30%（成長型對照無依據）", "lo", "hi", "7.13× → 4.99×／9.27×"),
    ("evs_si", "SI 其他 EV／NTM 營收（區間上緣；下緣 5×＝基準）", None, "hi", "5× → 8×"),
    ("mult_si_adj", "自用 AI 營運成本加回倍數（讀法 2）", "lo", "hi", "中位數 × 0.7／× 1.4"),
    ("he_all", "Tokenomics IF_HoldEcon 四世代同時", "lo", "hi", "低成本／高成本（k0 隨之重新校準、FY2030 k 固定 → 持有成本同時推動營收）"),
    ("tk_all", "Tokenomics 一致低／高成本（持有成本、營運費用、IT 與機房資本支出、IT 折舊同動）", "lo", "hi", "Tokenomics 低成本／高成本欄"),
    ("conv_k", "k 收斂年數", "lo", "hi", "4 → 2／6 年"),
    ("k_flat", "k 不收斂（三情境 k 維持 2026Q3 校準值 k0）", None, "hi", "conv_k 極大"),
    ("tax_tv", "終值期稅率 12% 永續（舊口徑）", "lo", None, "21% → 12%"),
    ("tv_alt", "終值改用 FY2032 自由現金流", None, "hi", "價值驅動 → UFCF ×(1＋g)"),
    ("gw_ye25", "2025 年底名目算力（汰換時點）", "lo", "hi", "0.8 → 0.7／0.9 GW"),
    ("gw_ye24", "2024 年底名目算力（汰換時點）", "lo", "hi", "0.15 → 0.1／0.25 GW"),
    ("life_fac", "機房折舊年限", "lo", "hi", "20 → 15／25 年"),
    ("phi_0", "資本支出當期折舊比例", "lo", "hi", "0.25 → 0／0.5"),
    ("gl_ads", "X 廣告長期成長", "lo", "hi", "3% → −3%／8%"),
    ("g_nl", "客戶發射次數年增", "lo", "hi", "5% → 0／10%"),
    ("g_rpl", "每次客戶發射收入年增", "lo", "hi", "3% → 0／6%"),
    ("gl_ld", "發射與開發長期成長", "lo", "hi", "8% → 0／15%"),
    ("dil", "股數年稀釋", "lo", "hi", "1% → 0／2%"),
    ("wc_k", "營運資金 ÷ 營收增量", "lo", "hi", "2.8% → 0／10%"),
    ("evs_space", "Space EV／營收", "lo", "hi", "3× → 2×／6×"),
    ("w_dcf", "DCF 權重", "lo", "hi", "50% → 30%／70%"),
    ("dil_x", "評價股數加入潛在稀釋（績效獎酬、選擇權、RSU、EchoStar）", None, "hi", "13.58 → 15.82 十億股"),
    ("sbc_add", "DCF 加回股權報酬", None, "hi", "不加回 → 加回（稅後）"),
]
def _tkset(case):
    out = []
    for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3"), ("VR200", "vr")):
        out += [(f"he_{nm}", tk("IF_HoldEcon", code, case)), (f"opx_{nm}", tk("IF_OpexGW", code, case))]
    out += [("cit_b3", tk("IF_CapexIT", "GB300", case)), ("cit_vr", tk("IF_CapexIT", "VR200", case)), ("cfac", tk("IF_CapexFacility", "VR200", case))]
    out += [(f"dpi_{nm}", tk("IF_DeprIT", code, case)) for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3"))]
    return out
SE_SPECIAL = {"he_all|低": [(f"he_{nm}", tk("IF_HoldEcon", code, "低成本")) for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3"), ("VR200", "vr"))],
              "he_all|高": [(f"he_{nm}", tk("IF_HoldEcon", code, "高成本")) for code, nm in (("H100", "h"), ("GB200", "b2"), ("GB300", "b3"), ("VR200", "vr"))],
              "tk_all|低": _tkset("低成本"), "tk_all|高": _tkset("高成本"), "k_flat|高": [("conv_k", 1000000)]}
SE["A1"] = "Sens — 一次改一個輸入（In 頁敏感度低／高值），整本活頁簿以 LibreOffice 重算後的加權目標價（Val 頁「目標價（DCF 與 Comparable 加權）」）；由 spacex/make_v01.sh 產生，建置時寫入（數值）"
SE["A2"] = "相對基準＝情境目標價 ÷ 基準加權目標價 − 1；龍捲風依絕對值排序（HTML）。基準與各案例的其他輸入相同"
hdr(SE, 3, ["驅動", "設定", "方向", "加權目標價", "相對基準", "鍵", "說明", "收支關卡 B（五年）", "FY2028 S2 比值（B）"])
r = 4
SE_ROWS = []
for key, lab, lo, hi, note in SENS_CASES:
    for side, val in (("低", lo), ("高", hi)):
        if val is None: continue
        SE.cell(r, 1, f"{lab}：{side}"); SE.cell(r, 2, note); SE.cell(r, 3, side); SE.cell(r, 6, key)
        res = SENS.get(f"{key}|{side}")
        if res is not None:
            c = SE.cell(r, 4, res["tp"]); c.number_format = "#,##0.00"; c.font = BLUE
            SE.cell(r, 5, f"=D{r}/Val!$B${VR_['tp']}-1").number_format = "+0.0%;-0.0%;0.0%"
            for cc, kk in ((8, "gateB"), (9, "s2B")):
                if kk in res:
                    c = SE.cell(r, cc, res[kk]); c.number_format = "#,##0.0" if kk == "gateB" else "0.00"; c.font = BLUE
        SE.cell(r, 7, "重算值（make_v01.sh）")
        SE_ROWS.append((key, side, val, r)); r += 1
SE_END = r - 1
r += 1
SE.cell(r, 1, "基準加權目標價（Val 頁）"); SE.cell(r, 4, f"=Val!B{VR_['tp']}").number_format = "#,##0.00"; SE_BASE = r
SE.cell(r, 8, f"=CF!C{GATE}").number_format = "#,##0.0"; SE.cell(r, 9, f"=Seg!I{SR['B']['s2']}").number_format = "0.00"; r += 1
SE.cell(r, 1, "敏感度重算時的基準值（應與上一列相同；不同表示敏感度已過期）")
if "base|base" in SENS:
    c = SE.cell(r, 4, SENS["base|base"]["tp"]); c.number_format = "#,##0.00"; c.font = BLUE
    SE.cell(r, 5, f"=D{r}-D{r - 1}").number_format = "0.000"
SE_STALE = r; r += 1
SE.column_dimensions["A"].width = 52; SE.column_dimensions["B"].width = 40; SE.column_dimensions["G"].width = 24

# =====================================================================================================
# Sources
SO["A1"] = "Sources — 來源（編號、文件、日期、網址、標記）"
hdr(SO, 3, ["編號", "文件", "日期", "網址", "標記", "用途"])
SRCS = [
    ("SpaceX 10-Q 2026Q3（Andy 2026-10-30 提供原檔：mag/extract/docs/2026Q3/spacex_2026Q3_10-Q.pdf）", "2026-10-22", "https://ir.spacex.com（SEC 網站本環境不可達）", "Verified",
     "分部、收入拆分、合併損益、現金流、資產負債、債務、承諾、客戶集中、名目算力定義（事實總帳 data/spacex_facts_2026q3.json）"),
    ("SpaceX 2026Q3 財報新聞稿（8-K 附件；mag/extract/docs/2026Q3/spacex_2026Q3_新聞稿.pdf）", "2026-10-20", "https://ir.spacex.com", "Interested-party", "2026Q2 比較欄、KPI（發射、訂戶、ARPU、名目算力）"),
    ("SpaceX 2026Q2 財報新聞稿（WebFetch；PDF 未入庫）", "2026-07", "https://ir.spacex.com", "Interested-party", "2026Q1、2025Q2、1H、Backlog 6-30、Q2 新簽雲端合約 14.1"),
    ("媒體轉述 SpaceX S-1／424B4（Motley Fool、Yahoo Finance、Implicator、Beancount）", "2026-05～07", "見事實總帳 seg.*.fy2025 的 url", "二手轉述 S-1", "FY2025 分部營收與營業利益（Space 資本支出未取得）"),
    ("Tokenomics 快照 spacex/data/tokenomics_snapshot_v21.7.json（自 alphabet/data/ 複製；20261025_Tokenomics_v21.7.xlsx）", "2026-10-29", "repo 內（YinchenChang/Tokenomics）", "Derived",
     "IF_HoldEcon、IF_OpexGW、IF_CapexIT／Facility、IF_DeprIT、IF_FacilityGW"),
    ("MarketBeat SPCX 預測頁", "2026-10-29", "https://www.marketbeat.com/stocks/NASDAQ/SPCX/forecast/", "Interested-party／Verified", "現價 143.89（4:00 PM 收盤）、共識目標價 244.97（46 位）"),
    ("Equibles SPCX 摘要", "2026-10-29", "https://equibles.com/stocks/spcx/summary", "Verified", "10-29 收盤 143.89（第二來源）"),
    ("Digrin SPCX 月線", "2026-10-30", "https://www.digrin.com/stocks/detail/SPCX/price", "Derived", "9 月收盤 150.86（評價日股價，單一來源）"),
    ("美國財政部每日殖利率曲線（10 Yr）", "2026-10-27", "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?type=daily_treasury_yield_curve&field_tdr_date_value_month=202610", "Verified", "無風險利率 4.99%"),
    ("同業檔 spacex/data/peers_si_20261030.json、peers_connectivity_20261030.json（MarketScreener）", "2026-10-08／10-23", "repo 內", "Derived", "SOTP 倍數"),
    ("共識檔 spacex/data/consensus_spcx_20261030.json（MarketBeat、Simply Wall St 過期、Shacknews）", "2026-10-30", "repo 內", "Interested-party", "共識目標價、過期年度共識"),
    ("A 段事實總帳報告 spacex/docs/reports/20261030_spacex_v0.1_A段_事實總帳.md", "2026-10-30", "repo 內", "—", "名目算力口徑、Valor、找不到的項目"),
]
for i, (doc, dt, url, tag, use) in enumerate(SRCS, 1):
    for j, v in enumerate((f"S{i:02d}", doc, dt, url, tag, use), 1): SO.cell(3 + i, j, v)
for col, w in zip("ABCDEF", [6, 80, 12, 70, 18, 60]): SO.column_dimensions[col].width = w

# =====================================================================================================
# README
RD["A1"] = "SpaceX（Nasdaq: SPCX，含 xAI）個股模型 v0.1（2026-10-30；2026Q3）— 命題模組＋官方營收分部 → 損益表"
RD["A3"] = ("命題：SI 分部的算力投入（9M26 SI 資本支出約 400 億美元、名目算力 1.7 GW）能否由 SI 營收（對外算力出租＋Grok／Cursor／X）回本；"
            "Connectivity（Starlink）現金流與 IPO 資金能支撐多久；加計 AI 折舊與 Cursor 收購獎酬後的每股價值是否高於現價。")
RD["A4"] = "核心問題：FY2028 SI 每 GW 年營收 ÷ Tokenomics 每 GW 年持有成本（與 OpenAI PxQ 頁同口徑）；收支關卡＝FY2026–FY2030 自由現金流合計（扣融資租賃本金與以融資取得 PP&E；Valor 失敗售後租回本金列融資，不扣）。"
VB = lambda k, c="B": f"Val!{c}{VR_[k]}"
PB_ = PR_["B"]
RD["A5"] = f'="投資建議："&{VB("rec")}&"（加權目標價 US$"&TEXT({VB("tp")},"0.00")&"，上檔空間 "&TEXT({VB("up")},"+0.0%;-0.0%")&"；現價 "&TEXT(px,"0.00")&"，2026-10-29 收盤）"'
RD["A6"] = (f'="重點觀察 1：目標價 "&TEXT({VB("tp")},"0.0")&"＝DCF "&TEXT({VB("tpdcf")},"0.0")&"、分部加總 "&TEXT({VB("tpcomp")},"0.0")&" 各半；情境 A／B／C 為 "'
            f'&TEXT({VB("scen")},"0.0")&"／"&TEXT({VB("scen", "C")},"0.0")&"／"&TEXT({VB("scen", "D")},"0.0")&"；市場共識目標價 "&TEXT(cons_tp,"0")&"，現價隱含 NTM EV/EBITDA "&TEXT({VB("mkt_imp")},"0")&" 倍。"')
RD["A7"] = (f'="重點觀察 2：S2 核心問題比值（FY2028 SI 每 GW 年營收 ÷ Tokenomics 持有成本）A／B／C＝"&TEXT({VB("s2")},"0.00")&"／"&TEXT({VB("s2", "C")},"0.00")&"／"&TEXT({VB("s2", "D")},"0.00")'
            f'&"（2026Q3 校準 "&TEXT(k0,"0.00")&"）；SI 名目算力 FY2028 年底 "&TEXT(Thesis!I{T["end"]},"0.0")&" GW，FY2027 SI 投資 "&TEXT(Thesis!H{T["si_inv"]},"0")&"。"')
RD["A8"] = (f'="重點觀察 3：收支關卡（FY2026–2030 自由現金流，收支關卡口徑）B "&TEXT(CF!C{GATE},"0.0")&"（公司口徑 "&TEXT(CF!C{GATE + 1},"0.0")&"），"&CF!C{GATE + 4}'
            f'&"；情境 B 合併營業利益率 FY2028 "&TEXT(PL!I{PB_["om"]},"0.0%")&"（新增 AI 折舊 "&TEXT(PL!I{PB_["da_ai"]},"0.0")&"）。"')
r = 10
RD.cell(r, 1, "驅動 → 推導對照表（每頁輸入、推導量、因果方向）").font = BOLD; r += 1
hdr(RD, r, ["頁", "輸入（In 頁鍵）", "推導量", "因果方向"]); r += 1
DRV = [("Thesis", "gw_0、si_cx_rr → pace、ext_sh、mix0_*、selfb、Tokenomics IF_HoldEcon／OpexGW／CapexIT／Facility", "SI 名目算力、對外／自用 GW、SI 投資、AI 折舊、自用增量成本、影子效益",
        "SI 投資 ÷ 每 GW 投資 → GW；GW → 投資 → 折舊 → SI 營業利益"),
       ("Seg", "k30_A／B／C、conv_k、g27_nr、gl_nr、gl_ads、nl_*、g_rpl、gl_ld、nadd_A／B／C、g_arpu、gl_eg", "三分部、六條子列、兩條 SI 命題子列、合併營收；非 AI 資本支出",
        "對外 GW × k × 持有成本；發射 × 每次收入；訂戶 × ARPU"),
       ("PL", "m_spl、m_cn、m_so、sbc_*、sbc_cur_left、da_*、tax", "分部調整後 EBITDA、分部營業利益、合併營業利益、淨利、EPS", "分部營收 → EBITDA − 折舊 − 股權報酬 → 營業利益 → EPS"),
       ("CF", "wc_k、cx_cn、cx_spl、fin_sh、mat_*、lease_*、spec_*、min_cash", "公司口徑 FCF、收支關卡口徑 FCF、新債、淨現金", "淨利＋折舊＋SBC − 資本支出 → FCF → 舉債"),
       ("Val", "rf、beta、erp、kd、g_tv、ronic、同業倍數、evs_space、tp_w*、w_dcf", "DCF、分部加總、目標價、上檔空間、投資建議", "無槓桿 FCF 折現＋分部 EBITDA × 倍數 → 情境加權 → 50／50"),
       ("Sens", "In 頁敏感度低／高值", "加權目標價變動", "一次改一個輸入、整本重算")]
for row in DRV:
    for j, v in enumerate(row, 1): RD.cell(r, j, v)
    r += 1
r += 1
RD.cell(r, 1, "本版要點（v0.1：首版）").font = BOLD; r += 1
for t in ("官方分部 Space、Connectivity、SI；子列＝10-Q 附註 3 收入拆分；SI S&I 之下以 Customer B 衍生 72% 拆「對外雲端服務」與「Grok／Cursor／X」命題子列（審查第 4 條）。",
          "SI 名目算力以實際值為錨（2025Q2 0.4 → 2026Q3 1.7 GW，IT 口徑），未來路徑由 SI 投資季速（含以融資取得 PP&E）÷ Tokenomics 每 GW 投資推算（審查第 3 條）。",
          "收支關卡口徑扣融資租賃本金與以融資取得 PP&E；Valor 失敗售後租回本金列融資、不扣（審查第 2 條 10-30 更正），並列公司口徑 FCF。",
          "歷史欄：Q 頁 2025Q2–2026Q3 與 9M；FY2025* 為二手轉述 S-1（Space 資本支出未取得）。Cursor 2026-08-14 併入，預測以備考全季為起點（審查第 8 條）。",
          "Tesla 合併評估（董事會特別委員會）不建模，列事件風險；EchoStar 頻譜按成本（價值中性），交割發行股數列 Sens。"):
    RD.cell(r, 1, t); r += 1
r += 1
RD.cell(r, 1, "工作表：README → In → Q → Thesis → Seg → PL → CF → Val → Recon → Sens → Sources。藍字＝輸入、黑字＝公式；黃底＝判斷題。建置：bash spacex/make_v01.sh").font = BOLD; r += 1
RD.cell(r, 1, "數字標記：Verified（經查核）／Interested-party（公司或賣方自述）／Analogy／Assumed（給區間）／Derived（換算）／二手轉述 S-1。本模型是研究整理，不構成投資建議。")
RD.column_dimensions["A"].width = 120; RD.column_dimensions["B"].width = 60; RD.column_dimensions["C"].width = 60; RD.column_dimensions["D"].width = 60
for rr in range(3, 9): RD.cell(rr, 1).alignment = Alignment(wrap_text=True, vertical="top")

if os.environ.get("ROWMAP"):
    json.dump({"T": T, "SR": SR, "PR": PR_, "CR": CR, "VR": VR_, "DC": DC, "GATE": GATE, "SUMROW": SUMROW, "SE_ROWS": SE_ROWS, "SE_BASE": SE_BASE,
               "SE_STALE": SE_STALE, "SE_SPECIAL": SE_SPECIAL, "IN_ROW": IN_ROW, "RCR": RCR, "QR": QR, "TH_HDR": TH_HDR},
              open(os.environ["ROWMAP"], "w", encoding="utf-8"), ensure_ascii=False, indent=0)
wb.save(A.out)
print("saved", A.out)
