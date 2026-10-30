# G2 校準（chat 端 10-30 二修；只供 Sens，不改基準）：解 k0（三情境同一 k0，FY2030 k 不變），使情境 B 的 2027Q1 合併營收（Recon 第 5 節幾何內插）＝指引中值 35。
# 每次試算整本 LibreOffice 重算；割線法。用法：python3 spacex/g2cal_v02.py <工作資料夾> <輸出 json>
import json, os, subprocess, sys, shutil, tempfile
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
WD, OUT = sys.argv[1], sys.argv[2]
os.makedirs(WD, exist_ok=True)
G2C = os.path.join(HERE, "data", "g2cal_v02_20261030.json")


def run(k0, i):
    raw = os.path.join(WD, f"g{i}.xlsx"); rmf = raw + ".rowmap.json"
    env = {**os.environ, "ROWMAP": rmf, "SPCX_NO_G2CAL": "1"}
    subprocess.run([sys.executable, os.path.join(HERE, "build_spacex.py"), raw, "--set", f"k0_ovr={k0}"], check=True, stdout=subprocess.DEVNULL, env=env)
    out = os.path.join(WD, "out"); prof = tempfile.mkdtemp(prefix="lo_g2_", dir=WD)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, raw],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env={**os.environ, "LC_ALL": "C.UTF-8"}, timeout=900)
    shutil.rmtree(prof, ignore_errors=True)
    rm = json.load(open(rmf, encoding="utf-8")); wb = openpyxl.load_workbook(os.path.join(out, os.path.basename(raw)), data_only=True)
    q1 = wb["Recon"].cell(rm["RCR"]["q1_m"], 3).value
    tgt = wb["In"].cell(rm["IN_ROW"]["g_rq1"], 5).value
    return q1, tgt, wb["Thesis"].cell(rm["T"]["k0"], 3).value


a, b = 1.8666, 4.0
fa, tgt, _ = run(a, 0); fb_, _, _ = run(b, 1)
fa -= tgt; fb_ -= tgt
for i in range(2, 12):
    c = b - fb_ * (b - a) / (fb_ - fa)
    fc, _, k0c = run(c, i); fc -= tgt
    print(i, round(c, 5), round(fc, 5), flush=True)
    a, fa, b, fb_ = b, fb_, c, fc
    if abs(fc) < 1e-4: break
json.dump({"k0": b, "q1_rev_B": fc + tgt, "target": tgt, "k0_check": k0c,
           "note": "情境 B 的 2027Q1 合併營收（Q4 26 與 FY2027 季均幾何內插）＝指引中值；三情境同一 k0；FY2030 k 不變"}, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("k0", b)
