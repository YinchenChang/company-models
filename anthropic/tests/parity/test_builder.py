"""builder 與資料分層的結構性測試（不需 LibreOffice）：E6、穩定 ID、OAI_Link 隔離、SRC／Inputs 與 yaml 一致、Derived 列公式化。"""
import json
import re
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "builder"))
from build import DERIVED, CALC_SHEETS  # noqa: E402
from common import e6_violations  # noqa: E402
from engine import Engine, current_model_path  # noqa: E402

SRC = yaml.safe_load((REPO / "data" / "anthropic_src.yaml").read_text(encoding="utf-8"))["rows"]
INP = yaml.safe_load((REPO / "data" / "anthropic_inputs.yaml").read_text(encoding="utf-8"))["rows"]
REG = json.loads((REPO / "builder" / "id_registry.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(current_model_path())


@pytest.fixture(scope="module")
def eng():
    return Engine(current_model_path())


def _same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-12 * max(abs(a), abs(b), 1)
    return a == b


def test_e6_no_constants_in_calc_sheets(wb):
    """E6：計算頁（A2：Demand、Revenue；A3：Compute、Cost）公式不得內含常數；恆等式只容許 1−比例、1＋成長率、年數 +1；定義常數在 Inputs。"""
    present = [s for s in CALC_SHEETS if s in wb.sheetnames]
    assert {"Demand", "Revenue", "Compute", "Cost", "Funding", "Reverse"} <= set(present)
    bad = [c for s in present for c in e6_violations(wb[s])]
    assert bad == [], bad[:10]
    n = sum(1 for s in present for row in wb[s].iter_rows(min_row=5) for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert n >= 800


def test_oai_link_not_referenced_by_calc_sheets(wb):
    """D20：OAI_Link 只被 Checks 引用。"""
    for ws in wb.worksheets:
        if ws.title in ("OAI_Link", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "OAI_" not in c.value, f"{ws.title}!{c.coordinate} 引用 OAI_Link：{c.value}"


def test_stable_ids():
    """SRC_ANT_nnn 與 yaml 一致且不重複；INP_nnn、Cnn 由 registry 發出、不重用退役號碼。"""
    ids = [r["id"] for r in SRC]
    assert len(ids) == len(set(ids))
    assert ids == sorted(ids) and ids[-1] == f"SRC_ANT_{len(ids):03d}"
    assert all(REG["SRC"][r["key"]] == r["id"] for r in SRC)
    for kind in ("INP", "CHK"):
        vals = list(REG[kind].values())
        assert len(vals) == len(set(vals))
        assert not set(vals) & set(REG.get(f"retired_{kind}", []))
    assert {r["key"] for r in INP} == set(REG["INP"])
    assert set(REG.get("retired_INP", [])) >= {"INP_016", "INP_067", "INP_068", "INP_069"}     # A2：每任務 token 改取 Tokenomics；A3：容量上限佔位改接 Compute


def test_src_and_inputs_values_equal_yaml(eng):
    """SRC_ANT／Inputs 每個有值的列：Excel 具名範圍的值等於 yaml（Derived 公式列除外，另測）。"""
    for r in SRC:
        if r["value"] is not None and not DERIVED.get(r["key"]):
            assert _same(eng.get_name(r["id"]), r["value"]), r["id"]
    for r in INP:
        assert _same(eng.get_name(REG["INP"][r["key"]]), r["value"]), r["key"]


def test_derived_rows_are_formulas(wb):
    """yaml 檔頭：標記 Derived 的 SRC 列改為公式（或保留報導值並註明）；公式只引用 SRC_ANT／INP 具名範圍。"""
    ws = wb["SRC_ANT"]
    rows = {ws.cell(r, 2).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value}
    for r in SRC:
        if r["tag"] != "Derived":
            continue
        n = rows[r["key"]]
        f = ws[f"E{n}"].value
        if DERIVED[r["key"]]:
            assert isinstance(f, str) and f.startswith("=") and re.fullmatch(r"[=+\-*/()A-Z_0-9.1]+", f.replace("SRC_ANT_", "S").replace("INP_", "I")), (r["id"], f)
            assert ws[f"V{n}"].value == "Derived→公式"
        else:
            assert ws[f"V{n}"].value.startswith("Derived→保留報導值")


def test_inputs_ranges_and_tags():
    """共同規則第 4 節：Analogy／Assumed 一律給區間；低 ≤ 值 ≤ 高；標記只限 Assumed／Analogy／Decision。"""
    for r in INP:
        assert r["tag"] in ("Assumed", "Analogy", "Decision"), r["key"]
        assert r["lo"] is not None and r["hi"] is not None, r["key"]
        if isinstance(r["value"], str):          # 文字鍵（Tokenomics 任務名稱）：Decision，低＝高＝值
            assert r["tag"] == "Decision" and r["lo"] == r["hi"] == r["value"], r["key"]
        else:
            assert r["lo"] <= r["value"] <= r["hi"], r["key"]
        assert r.get("basis"), r["key"]


def test_named_outputs_for_a3(eng):
    """A3／A4 會用到的具名範圍（與 OpenAI v0.6 同名者）存在且為 6 年。"""
    for n in ("DEM_Tok_Paid_Top", "DEM_Tok_Paid_Mid", "DEM_Tok_Paid_Low", "DEM_Tok_Free_Top", "DEM_Tok_Free_Mid", "DEM_Tok_Free_Low",
              "DEM_Tok_Paid", "DEM_Tok_Free", "DEM_Tok_Total", "DEM_Tok_API", "REV_Sub", "REV_API", "REV_Gross", "REV_PartnerShare", "REV_Net",
              "REV_GrossCapped", "REV_NetCapped", "REV_CapFactor", "REV_Consumer", "REV_Enterprise"):
        v = eng.get_name(n)
        assert isinstance(v, list) and len(v) == 6 and all(isinstance(x, (int, float)) for x in v), n
    assert eng.get_name("CHK_Errors") == 0


def test_named_outputs_for_a4(eng):
    """A3 輸出、A4 會用到的具名範圍（與 OpenAI v0.6 同名）存在且為 6 年；容量上限 2025、2026＝1；Revenue 與 Compute 一致。"""
    for n in ("CMP_SupplyGW", "CMP_Supply_VReq", "CMP_InfGW_Eff", "CMP_InfGW", "CMP_RDGW", "CMP_DemandGW", "CMP_InfAvailGW", "CMP_CapFactor", "CMP_Eta",
              "COST_Compute", "COST_ContractPay", "COST_OwnedCapex", "COST_NonCompExSBC", "COST_SBC", "COST_FullCash", "COST_PropRev_VR",
              "COST_PropFull_VR", "COST_PropGap_VR", "COST_Coverage", "COST_CloudGM", "COST_ComputeEcon", "COST_GapCash"):
        v = eng.get_name(n)
        assert isinstance(v, list) and len(v) == 6 and all(isinstance(x, (int, float)) for x in v), n
    cap = eng.get_name("CMP_CapFactor")
    assert cap[0] == 1 and cap[1] == 1
    assert eng.get_name("REV_CapFactor") == cap
    assert abs(eng.get_name("COST_Gap2025")) < 1e-9                     # 2025 算力成本＝說明書 7.33


def test_reverse_not_fed_back(wb):
    """D19：Reverse（管理層目標反解）只作對照，不回饋基準——除 Reverse 本頁與 Checks 外，任何公式不得引用 Reverse! 或 RVS_ 名稱。"""
    assert "Reverse" in wb.sheetnames
    for ws in wb.worksheets:
        if ws.title in ("Reverse", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "Reverse!" not in c.value and "RVS_" not in c.value, f"{ws.title}!{c.coordinate} 引用 Reverse：{c.value}"


def test_named_outputs_v01(eng):
    """v0.1（A4）命題 2 與反向模式輸出（與 OpenAI v0.6 同名）存在；2025 為實際年（流量欄空白），2026–2030 為數值；外部資金需求非負、年底現金 ≥ 最低現金。"""
    for n in ("FND_FCF", "FND_Committed", "FND_EquityIn", "FND_MinCash", "FND_CashEnd", "RVS_Target", "RVS_MultAPI", "RVS_MultSub", "RVS_MultProp", "RVS_FCF",
              "COST_PropGap_GW", "COST_ComputeAnchor", "COST_PropGap_VR_Anchor", "COST_LeaseRent", "COST_PropGap_VR_Rent", "COST_PropGap_VR_D22"):
        v = eng.get_name(n)
        assert isinstance(v, list) and len(v) == 6 and all(isinstance(x, (int, float)) for x in v), n
    for n in ("FND_ExtNeed", "FND_ExtNeedCum", "FND_CashOpen", "FND_Headroom", "RVS_ExtNeedCum", "FND_ExtNeedCumIPO", "FND_ExtNeedCumCond", "FND_ExtNeedCumAll",
              "FND_ExtNeedCumAnchor", "FND_ExtNeedCumRent", "FND_ExtNeedCumD22", "FND_ExtNeedCumAdverse"):
        v = eng.get_name(n)
        assert len(v) == 6 and all(isinstance(x, (int, float)) for x in v[1:]), n
    ext, end, mn = eng.get_name("FND_ExtNeed"), eng.get_name("FND_CashEnd"), eng.get_name("FND_MinCash")
    assert all(x >= 0 for x in ext[1:])
    assert all(e >= m - 1e-9 for e, m in zip(end[1:], mn[1:]))
    for n in ("FND_FirstGapYear", "FND_ExtNeedPeak", "FND_PeakYear", "FND_Contingent", "FND_CommitAfter2030", "COST_PartnerGapMax"):
        assert eng.get_name(n) not in (None, ""), n
    assert eng.get_name("CHK_Errors") == 0


# ── v0.1.3 B：Bridge、EtaDiag（唯讀診斷頁）─────────────────────────────────
def test_bridge_not_fed_back(wb):
    """Bridge、EtaDiag 只作對照與診斷：除本頁與 Checks 外，任何公式不得引用 Bridge!、EtaDiag! 或 ETA_ 名稱。"""
    assert "Bridge" in wb.sheetnames and "EtaDiag" in wb.sheetnames
    assert wb.sheetnames.index("Bridge") == wb.sheetnames.index("Cost") + 1       # 放在 Cost 之後
    for ws in wb.worksheets:
        if ws.title in ("Bridge", "EtaDiag", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert not re.search(r"\b(Bridge|EtaDiag)!|\bETA_", c.value), f"{ws.title}!{c.coordinate} 引用診斷頁：{c.value}"


def test_eta_decomposition_does_not_change_eta(eng):
    """η 拆解只作診斷：η 等於 Compute 的值（v0.1.4 起隨 TK v19.1 為 0.525427732668394；v0.1.3 為 0.52729641508883，差異全來自 TK IF_TokGW_Luna）；乘積 × 殘差 ＝ η；殘差為正。"""
    eta = eng.get_name("CMP_Eta2025")
    assert abs(eta - 0.525427732668394) < 1e-9
    assert abs(eng.get_name("ETA_Eta") - eta) < 1e-12
    prod = eng.get_name("ETA_fP") * eng.get_name("ETA_fU") * eng.get_name("ETA_fW")
    assert abs(prod - eng.get_name("ETA_Prod")) < 1e-12
    assert abs(prod * eng.get_name("ETA_Resid") - eta) < 1e-12
    assert eng.get_name("ETA_Resid") > 0
    assert 0 < eng.get_name("ETA_fU") < 2 and eng.get_name("ETA_fW") > 0 and eng.get_name("ETA_fP") > 0


def test_eta_prod_derate_section(eng):
    """v0.1.4：f_U'＝f_U × TK 生產折減；殘差' ＝ η ÷ (f_P × f_U' × f_W)；η 不變，殘差 ×（1 ÷ 生產折減）。"""
    eta = eng.get_name("CMP_Eta2025")
    pd = eng.get_name("TK_IF_ProdDerate")
    assert abs(eng.get_name("ETA_fU_Prod") - eng.get_name("ETA_fU") * pd) < 1e-12
    prod2 = eng.get_name("ETA_fP") * eng.get_name("ETA_fU_Prod") * eng.get_name("ETA_fW")
    assert abs(prod2 - eng.get_name("ETA_Prod2")) < 1e-12
    assert abs(prod2 * eng.get_name("ETA_Resid2") - eta) < 1e-12
    assert abs(eng.get_name("ETA_Resid2") - eng.get_name("ETA_Resid") / pd) < 1e-12


def test_bridge_ties_to_calc_sheets(eng):
    """Bridge 逐層列的值與計算頁一致（核對列為 0）；有 Bridge 時 Checks e1_* 全 OK。"""
    rows = eng.get("Bridge", "A8:I140")
    byid = {r[0]: r for r in rows if r[0] != ""}
    for rid, r in byid.items():
        lab = str(r[1])
        if lab.startswith("核對"):
            assert all(abs(x) < 1e-9 for x in r[3:9] if isinstance(x, (int, float))), f"{rid} {lab}"
    gap = eng.get_name("COST_PropGap_VR")
    gap_row = next(r for r in byid.values() if str(r[1]).startswith("每 VR 等值 GW：差額（含股權報酬；命題 1）"))
    assert all(abs(a - b) < 1e-9 for a, b in zip(gap_row[3:9], gap))
    sup_row = next(r for r in byid.values() if r[1] == "供給 GW 合計")
    assert all(abs(a - b) < 1e-9 for a, b in zip(sup_row[3:9], eng.get_name("CMP_SupplyGW")))


def _s1f_rows(wb):
    """Cost 第二十四節（v0.1.6 正式 S-1 對照）的列號範圍：節標題至頁尾。"""
    ws = wb["Cost"]
    start = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("二十四、"))
    return ws, set(range(start, ws.max_row + 1))


def test_s1f_section_not_fed_back(wb):
    """v0.1.6 W3：Cost 第二十四節只對照（不回饋基準）——除本節與 Checks 外，任何公式不得引用本節儲存格或本節具名範圍。"""
    ws, rows = _s1f_rows(wb)
    names = [n for n, v in wb.defined_names.items() if v.attr_text.startswith("Cost!")
             and int(re.search(r"\$(\d+)", v.attr_text).group(1)) in rows]
    assert "COST_Eta2026H1" in names and "COST_PropGap_VR_EtaH1" in names
    for sh in wb.worksheets:
        if sh.title == "Checks":
            continue
        for row in sh.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")) or (sh.title == "Cost" and c.row in rows):
                    continue
                refs = [int(x) for x in re.findall(r"Cost!\$?[A-Z]{1,3}\$?(\d+)", v)]
                if sh.title == "Cost":
                    refs += [int(x) for x in re.findall(r"(?<![A-Za-z!_$])\$?[D-I]\$?(\d+)\b", re.sub(r"[A-Za-z_]+![$A-Z0-9:]+", "", v))]
                assert not [r for r in refs if r in rows], f"{sh.title}!{c.coordinate} 引用第二十四節：{v}"
                assert not [n for n in names if re.search(rf"\b{n}\b", v)], f"{sh.title}!{c.coordinate} 引用第二十四節具名範圍：{v}"


def test_s1f_w2_and_eta_unchanged(eng):
    """v0.1.6 W2：2026 營收基準＝正式 S-1 上半年＋Q3＋Q4（Q3 持平）；基準 η 不因 W3 改變；B 情境 2025 欄＝基準。"""
    s = {r["key"]: r["id"] for r in SRC}
    i = {r["key"]: REG["INP"][r["key"]] for r in INP}
    assert eng.get_name(i["s1f_switch"]) == 1 and eng.get_name(i["u1_switch"]) == 1
    tgt = eng.get_name(s["rev_s1f_2026h1"]) + eng.get_name(s["rev_2026_q3"]) * (2 + eng.get_name(i["q4_growth"]))
    assert abs(eng.get_name("REV_Target")[1] - tgt) < 1e-9
    assert abs(eng.get_name("CMP_Eta2025") - 0.525427732668394) < 1e-9
    assert eng.get_name("COST_Eta2026H1") > 0
    for a, b in ((eng.get_name("COST_PropGap_VR_EtaH1")[0], eng.get_name("COST_PropGap_VR")[0]),
                 (eng.get_name("COST_Coverage_EtaH1")[0], eng.get_name("COST_Coverage")[0])):
        assert abs(a - b) < 1e-9
