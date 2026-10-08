"""Excel 優先（沿用 Tokenomics v5.7 精神）：藍字輸入由 Excel 擁有，只在仍為舊預設值時才寫入新預設。

隱藏頁 _Defaults 記錄「上一次 builder 寫入的預設值」。重建時，對每個輸入格（SRC_ANT E–G、Inputs E–G）：
  Excel 值 ＝ 舊預設 → 改寫為新預設（builder 的更新生效）；
  Excel 值 ≠ 舊預設 → 保留 Excel 值並記入 restore_log（Excel 值保留數）。
TK_Link 的值由 builder 從 Tokenomics 快照重寫（更新快照＝本模型的修補版），不在此保護。
"""
from __future__ import annotations

import openpyxl

SHEET = "_Defaults"
COLS = {"SRC_ANT": ("E", "F", "G"), "Inputs": ("E", "F", "G")}   # 藍字輸入欄（值、低、高）


def read_defaults(path):
    wb = openpyxl.load_workbook(path)
    cur = {}
    for sh in COLS:
        if sh not in wb.sheetnames:
            continue
        ws = wb[sh]
        for row in ws.iter_rows(min_row=5):
            rid = row[0].value
            if not rid:
                continue
            for col in COLS[sh]:
                cur[(sh, rid, col)] = ws[f"{col}{row[0].row}"].value
    old = {}
    if SHEET in wb.sheetnames:
        for r in wb[SHEET].iter_rows(min_row=2, values_only=True):
            if r[0]:
                old[(r[0], r[1], r[2])] = r[3]
    return cur, old


def write_defaults(wb, defaults: dict):
    ws = wb.create_sheet(SHEET)
    ws.append(["sheet", "id", "col", "builder 預設值（上次寫入）"])
    for (sh, rid, col), v in defaults.items():
        ws.append([sh, rid, col, v])
    ws.sheet_state = "hidden"


def merge(new_defaults: dict, base_path, log_path=None):
    """回傳 (最終值, matched, kept, unmatched_ids)。"""
    cur, old = read_defaults(base_path)
    final, kept, matched, lines = dict(new_defaults), 0, 0, []
    for key, newv in new_defaults.items():
        if key in cur and key in old:
            matched += 1
            if cur[key] != old[key]:        # Excel 已改過
                final[key] = cur[key]
                kept += 1
                lines.append(f"{key}: code default {newv!r} (previous {old[key]!r}) -> Excel {cur[key]!r} kept")
    unmatched = [k for k in cur if k not in new_defaults]
    if log_path:
        open(log_path, "w", encoding="utf-8").write(
            f"matched: {matched}; Excel value kept over code default: {kept}; unmatched: {len(unmatched)}\n" + "\n".join(lines)
            + "\n-- unmatched --\n" + "\n".join(map(str, unmatched)))
    return final, matched, kept, unmatched
