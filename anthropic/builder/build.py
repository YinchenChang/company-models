#!/usr/bin/env python3
"""產生 Anthropic 收支模型 Excel（v0.1；分段 A2→A4）。

用法：python3 builder/build.py --tk-dir <Tokenomics checkout> --oai-xlsx <OpenAI v0.6 xlsx> --out model/YYYYMMDD_Anthropic_v0.1.xlsx [--base <前一版>] [--date YYYY-MM-DD]
  --base：前一版 xlsx；有給時依 Excel 優先規則保留 Excel 內已改過的藍字輸入（builder/preserve.py）。
輸出經 LibreOffice 重算後才寫到 --out（內含快取值；engine 仍會強制重算）。紀錄檔寫到 build_out/（restore_log.txt、build_summary.json）。

分段設計：共用頁（README、SRC_ANT、TK_Link、OAI_Link、Inputs、Checks、_Defaults）由本檔產生；各段的計算頁由同目錄的模組產生，
依序載入存在者：a2.py（Demand、Revenue）、a3.py（Compute、Cost）、a4.py（Funding、Reverse）。每個模組提供
  build(ctx)   → 建立該段的工作表（ctx.wb）並把列位置存入 ctx.P[模組名]
  checks(ctx)  → 回傳該段的檢查列 [(穩定鍵, 檢查項, 公式, 期望, 種類, 說明)]；種類 eq／tol／warn／info
Python 只產生結構（公式文字）；參數一律來自 data/*.yaml 或 Tokenomics 快照（CLAUDE.md 第 1 節）。
"""
from __future__ import annotations

import argparse
import datetime as dt
import importlib
import json
import re
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import preserve  # noqa: E402
import source_rules  # noqa: E402
from common import F_BOLD, F_CALC, F_IN, F_NOTE, e6_violations, header, lo_recalc, nm, put, title  # noqa: E402
from oai_link import read_oai  # noqa: E402
from tk_link import NONNV_STATUS, PENDING_NOTE, read_nonnv, read_snapshot  # noqa: E402

VERSION = "v0.1"
STAGE = "v0.1"
REGISTRY_PATH = HERE / "id_registry.json"
SRC_YAML = ROOT / "data" / "anthropic_src.yaml"
INP_YAML = ROOT / "data" / "anthropic_inputs.yaml"
STAGE_MODULES = ("a2", "a3", "a4")
CALC_SHEETS = ("Demand", "Revenue", "Compute", "Cost", "Funding", "Reverse")      # E6 掃描與 OAI_Link 隔離的對象
SRC_COLS = ["SRC_ID", "key", "分組", "項目", "值", "低", "高", "單位", "期間", "日期（序列值）", "口徑", "來源", "網址", "文件日期", "擷取日", "標記",
            "利害方與誘因", "來源等級", "立場", "立場說明", "一手／二手", "處理狀態", "備註",
            "缺值（公式）", "缺出處（公式）", "區間順序異常（公式）", "Derived 未處理（公式）", "四欄標註"]
INP_COLS = ["INP_ID", "參數", "索引", "單位", "值", "低", "高", "標記", "依據", "區間理由", "用途／備註",
            "區間順序異常（公式）", "Analogy／Assumed 缺區間（公式）", "key"]
SRC_FIRST, INP_FIRST, TK_FIRST, TK_VCOL0 = 5, 5, 10, 10
EXCEL_EPOCH = dt.date(1899, 12, 30)

# 規格與 yaml 檔頭：標記 Derived 的 SRC 列改為公式（由其他 SRC 列推得）；公式中的常數改引用 SRC 新列或 Inputs 定義常數。
# 值為公式樣板：{SRC:key} 代 SRC 具名範圍、{INP:key} 代 Inputs 具名範圍；None＝保留報導值（報導者自行計算後公布的數字，視同觀測值）。
DERIVED = {
    "rev_2024_recognized_derived": "={SRC:rev_2025_recognized_b}/(1+{SRC:rev_2025_growth_reported})",
    "rev_2026_h1_derived": "={SRC:rev_2026_q1}+{SRC:rev_2026_q2}",
    "rev_partner_fee_2025": None,
    "rev_partner_fee_rate_2025": "={SRC:rev_partner_fee_2025}/({SRC:rev_2025_recognized_b}*{SRC:rev_mix_2025_cloud_partner_share})",
    "price_event_opus45_cut": "={SRC:price_opus45_in}/{SRC:price_opus4_in}-1",
    "cmp_spacex_total_b": "={SRC:cmp_spacex_monthly_fee}*{INP:spacex_term_months}",
    "cmp_spend_2025_derived_inference": "={SRC:cmp_spend_2025_actual_prospectus}-{SRC:cmp_spend_2025_training_proj}",
    "cmp_gw_target_nyt_2027": "={SRC:cmp_gw_target_nyt}*{SRC:cmp_gw_2027_vs_2026_multiple_nyt}",
    "cost_noncompute_opex_2025": "={SRC:cost_opex_2025}-{SRC:cost_compute_2025}",
    "cost_compute_to_rev_2025": "={SRC:cost_compute_2025}/{SRC:rev_2025_recognized}",
    "cost_employee_q1_2026_derived": "={SRC:cost_charity_match_q1_2026}/{SRC:cost_charity_match_share_employee_exp_q1_2026}",
    "strat_amzn_funded_cum_2026h1": "={SRC:strat_amzn_notes_cum_2025}+{SRC:strat_amzn_2026_series_g}+{SRC:strat_amzn_2026_series_h}",
}


def blank(v):
    return None if v == "" else v


def s(v):
    """yaml 的日期／整數期間一律轉文字寫入。"""
    if v is None:
        return None
    return v.isoformat() if isinstance(v, (dt.date, dt.datetime)) else str(v)


def serial(period):
    """期間為完整日期（YYYY-MM-DD）者轉 Excel 日期序列值（表示法轉換，非參數）；否則 None。"""
    t = s(period)
    if t and re.fullmatch(r"\d{4}-\d{2}-\d{2}", t):
        return (dt.date.fromisoformat(t) - EXCEL_EPOCH).days
    return None


# ═════════════════════ 資料與穩定 ID ═════════════════════
class Data:
    """SRC_ANT 與 Inputs 的資料列，以 key 取 ID 與列號。"""

    def __init__(self):
        reg = json.loads(REGISTRY_PATH.read_text(encoding="utf-8")) if REGISTRY_PATH.exists() else {}
        self.reg = reg
        for k in ("SRC", "INP", "CHK"):
            reg.setdefault(k, {})
        src = yaml.safe_load(SRC_YAML.read_text(encoding="utf-8"))
        self.src_rows = src["rows"]
        self.not_found = src.get("not_found", [])
        for r in self.src_rows:                       # SRC_ANT id 由 A1 發出、寫在 yaml；registry 只核對不改號
            old = reg["SRC"].get(r["key"])
            if old and old != r["id"]:
                sys.exit(f"SRC_ANT 穩定 ID 衝突：{r['key']} registry {old} ≠ yaml {r['id']}")
            if r["id"] in set(reg["SRC"].values()) - {old}:
                sys.exit(f"SRC_ANT ID 重複使用：{r['id']}")
            reg["SRC"][r["key"]] = r["id"]
        inp = yaml.safe_load(INP_YAML.read_text(encoding="utf-8"))
        self.inp_rows = inp["rows"]
        keys = [r["key"] for r in self.inp_rows]
        if len(keys) != len(set(keys)):
            sys.exit("Inputs key 重複")
        nxt = max([int(v[4:]) for v in reg["INP"].values()] + [int(x[4:]) for x in reg.get("retired_INP", [])] + [0]) + 1
        for r in self.inp_rows:
            if r["key"] not in reg["INP"]:
                reg["INP"][r["key"]] = f"INP_{nxt:03d}"
                nxt += 1
            r["id"] = reg["INP"][r["key"]]
        for k in [k for k in reg["INP"] if k not in set(keys)]:        # 被移除的 Inputs：號碼退役不重用
            reg.setdefault("retired_INP", []).append(reg["INP"].pop(k))
        self._src = {r["key"]: r for r in self.src_rows}
        self._inp = {r["key"]: r for r in self.inp_rows}
        self._src_row = {r["key"]: SRC_FIRST + i for i, r in enumerate(self.src_rows)}
        self._inp_row = {r["key"]: INP_FIRST + i for i, r in enumerate(self.inp_rows)}
        for k in DERIVED:
            if k not in self._src:
                sys.exit(f"DERIVED 指到不存在的 SRC key：{k}")
        for r in self.src_rows:
            if r["tag"] == "Derived" and r["key"] not in DERIVED:
                sys.exit(f"Derived 列未處理：{r['id']} {r['key']}")

    def save_registry(self):
        REGISTRY_PATH.write_text(json.dumps(self.reg, ensure_ascii=False, indent=0), encoding="utf-8")

    # 取 ID／列號
    def S(self, key):
        return self._src[key]["id"]

    def I(self, key):
        return self._inp[key]["id"]

    def src(self, key):
        return self._src[key]

    def inp(self, key):
        return self._inp[key]

    def src_row(self, key):
        return self._src_row[key]

    def inp_row(self, key):
        return self._inp_row[key]

    def has_inp(self, key):
        return key in self._inp

    def fill(self, tpl):
        """公式樣板：{SRC:key}、{INP:key} 代具名範圍。"""
        return re.sub(r"\{(SRC|INP):([A-Za-z0-9_]+)\}", lambda m: self.S(m.group(2)) if m.group(1) == "SRC" else self.I(m.group(2)), tpl)


class Ctx:
    def __init__(self, wb, D, snap, oai, date):
        self.wb, self.D, self.snap, self.oai, self.date = wb, D, snap, oai, date
        self.P = {}
        self.scan = {}


# ═════════════════════ 共用頁 ═════════════════════
def sheet_readme(ctx, summary):
    ws = ctx.wb.active
    ws.title = "README"
    title(ws, f"Anthropic 收支模型 {VERSION}（{STAGE}；{ctx.date}）",
          "命題（與 OpenAI v0.6 同一句）：Anthropic 每 VR 等值 GW 的年營收能否覆蓋每 GW 年全成本；若不能，缺口要多少外部資金、由誰以什麼條件提供。FY2025–FY2030，曆年制，2025 為實際校準年。")
    sn, oa = ctx.snap, ctx.oai
    lines = [
        ("本版範圍", f"{STAGE}：SRC_ANT、TK_Link（含 NonNV 讀表）、OAI_Link、Inputs、Demand（需求與 token 量）、Revenue（訂閱方案別、API 層級別、通路分成、淨額、容量上限）、Compute（加速器族、逐合約供給 GW、η、推論與研發 GW、容量上限、VR 等值）、"
                    "Cost（逐合約實付、自建資本支出、供應商持有成本、非算力成本、股權報酬、命題表；A4 補逐家對帳與敏感度）、Funding（命題 2）、Reverse（管理層目標反解，只作對照）、Checks。"),
        ("Compute", "加速器族 × 世代（NVIDIA Hopper／GB200／GB300／VR200；TPU v6e／v7；Trainium2／3；AMD MI455X）每 GW 年產能＝TK IF_TokGW_* ×（NonNV 產出比）；供給 GW 只由算力合約加總（D10 r1；機房租約只列對照）；"
                    "η＝2025 token 換算推論 GW ÷ 2025 推論支出換算 GW（D8）；研發 GW＝供給 ×（1 − 閒置）− 推論（D12）；容量上限 2025、2026 固定 1；VR 等值＝各族 Sol 產能比。"),
        ("Cost", "算力成本（現金）＝逐合約實付＋自建資本支出（D10、D11）；供應商持有成本＝GW × TK IF_HoldEcon × 持有比、雲端毛利；非算力成本 2025＝說明書營業費用 − 算力 − 平台抽成（D13、D5 r1），之後人數 × 每人成本；"
                 "股權報酬單列；命題表：每 VR 等值 GW 營收淨額、算力、非算力、全成本（含／不含股權報酬）、差額、覆蓋率。"),
        ("Funding", "自由現金流＝營收淨額（截頂後）− 算力成本（現金）− 非算力成本（不含股權報酬）；來源順序（D16 r1）：2025 年底現金 → 2026 已交割股權（Series G 30、Amazon G 5、Series H 65）→ 外部資金（補足至最低現金＝次年非算力 6 個月，D17）；"
                    "條件式（Amazon 15、Google 30、NVIDIA 10、AMD 5）與 IPO 約 100 只列情境；或有負債只列示；情境 S1–S7（IPO、條件式、逐家錨定招股書、機房租金、D22、合併）；回流對照（D18）。"),
        ("Reverse", "管理層營收目標（2026 18、2027 55、2028 70、2029 148；2030 無目標取正向）→ 只靠 API／只靠訂閱／兩線等比例的所需倍數 → 反向資金（同一支出）；只被 Checks 引用（D19）。"),
        ("Excel 為唯一計算引擎", "藍字＝輸入（Excel 擁有，重建時保留已改過的值）；黑字＝公式；綠字＝跨頁連結。builder 只產生結構（data/*.yaml → SRC_ANT／Inputs；Tokenomics → TK_Link；OpenAI → OAI_Link）。"),
        ("SRC_ANT", f"{summary['src']} 列：公司財務原始數據（A1 蒐集 376 列＋A2 新增 3 列供 Derived 列公式化）。標記 Derived 的 12 列：11 列改為公式、1 列（Reuters 自行計算後公布的通路費）保留報導值。"),
        ("TK_Link", f"Tokenomics {sn['version']}（{sn['file']}），master 提交 {sn['sha'][:7]}；{summary['tk_ok']} 個具名範圍（OpenAI v0.6 同一組 63 名＋A2 新增任務 token 5 名）＋ 讀表（非具名）{summary['tk_nnv']} 格：NonNV 18 格（規格 D6）與 PUE 3 格（A3；Tokenomics Inputs 頁）。"),
        ("OAI_Link", f"OpenAI v0.6 命題輸出快照（{oa['file']}；SHA-256 {oa['sha256'][:12]}…）；只被 Checks 引用（規格 D20）。"),
        ("Inputs", f"{summary['inp']} 列：假設（值、低、高、標記、依據、區間理由）；定義常數（年度、天數、單位換算）也在此（E6：公式不含常數）。"),
        ("Demand", "個人方案（Free／Pro／Max 5x／Max 20x）人數、企業席位（Team 標準／Premium、Enterprise）、任務類別（對話、程式代理、其他代理）× 每任務 token × 層級組合；API：2025 由 API 營收 ÷ 有效單價倒推、2026 由半校準總額倒推、2027 起任務成長 × 每任務 token 成長 × 價格彈性；輸出 DEM_Tok_*（層級 × 付費／免費）。"),
        ("Revenue", "API 牌價（價格事件依公告日天數加權）、有效單價（快取、批次、議價折扣）、訂閱方案別、API 層級別、其他、廣告（＝0，D3）、總額、雲端通路平台抽成（D5 r1）、淨額、個人 vs 企業；2025 校準（D14）、2026 半校準（D15 r1）、容量上限係數（＝Compute CMP_CapFactor；2025、2026 固定 1）。"),
        ("Checks", "C01 起；ERR 格數＝CHK_Errors（必須 0）；WARN 格數＝CHK_Warnings（只提示）。"),
        ("層級對應", "Haiku→低層（Tokenomics Luna）、Sonnet→中層（Sol）、Opus→頂層（Astra）；具名範圍沿用 OpenAI 的 Top／Mid／Low。"),
        ("標記", "Verified／Interested-party／Analogy／Assumed／Derived／Decision（Analogy、Assumed 一律附區間）。"),
        ("分層", "公司財務原始數據→SRC_ANT；AI 技術與算力→TK_Link（取自 Tokenomics）；假設→Inputs；OpenAI 對照→OAI_Link（不進計算）。"),
    ]
    for i, (a, b) in enumerate(lines):
        put(ws, f"A{4 + i}", a, F_BOLD)
        put(ws, f"B{4 + i}", b, F_CALC, wrap=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 130


def sheet_src(ctx, final):
    D, wb = ctx.D, ctx.wb
    ws = wb.create_sheet("SRC_ANT")
    title(ws, "SRC_ANT — 第 0 層 Source：Anthropic 公司財務原始數據（營收、價格、用戶、合約、成本、融資、管理層目標）",
          "只存已取得的原始訊息（非計算、非假設）。藍字＝原始值，由本活頁簿擁有。標記 Derived 的列：值欄改為公式（由其他 SRC_ANT 列推得），處理狀態欄說明；AI 技術與算力原始數據不在此登錄（取自 TK_Link）。",
          "來源等級（1 一手／2 二手已讀或轉載／未評）、立場、立場說明、一手／二手：依 builder/source_rules.py 的機械規則由 CC 初評（AB 欄）；利害方與誘因欄為 A1 逐列登錄。")
    header(ws, 4, SRC_COLS)
    for i, r in enumerate(D.src_rows):
        n = SRC_FIRST + i
        cl = source_rules.classify(r)
        if r["key"] in DERIVED:
            status = "Derived→公式" if DERIVED[r["key"]] else "Derived→保留報導值（報導者計算後公布）"
        else:
            status = "觀測"
        vals = [r["id"], r["key"], r["group"], r["metric"], None, None, None, r["unit"], s(r["period"]), serial(r["period"]), r.get("basis"),
                r["source"], r["url"], s(r["doc_date"]), s(r["retrieved"]), r["tag"], r.get("party"), cl["grade"], cl["stance"], cl["why"], cl["hand"],
                status, r.get("note")]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=n, column=c, value=v).font = F_CALC
        for col in "EFG":
            v = final[("SRC_ANT", r["id"], col)]
            put(ws, f"{col}{n}", v, F_IN if isinstance(v, (int, float)) else None)
        put(ws, f"X{n}", f'=IF(AND(COUNT($E{n}:$G{n})=0,LEFT($H{n},2)<>"定性",LEFT($H{n},2)<>"文字",LEFT($H{n},2)<>"日期"),1,0)')
        put(ws, f"Y{n}", f'=IF(OR($L{n}="",$M{n}=""),1,0)')
        put(ws, f"Z{n}", f'=IF(AND(ISNUMBER($E{n}),ISNUMBER($F{n})),IF($F{n}>$E{n},1,0),0)+IF(AND(ISNUMBER($E{n}),ISNUMBER($G{n})),IF($E{n}>$G{n},1,0),0)'
                         f'+IF(AND(ISNUMBER($F{n}),ISNUMBER($G{n})),IF($F{n}>$G{n},1,0),0)')
        put(ws, f"AA{n}", f'=IF(AND($P{n}="Derived",$V{n}="觀測"),1,0)')
        put(ws, f"AB{n}", "CC 初評", F_NOTE)
        if final[("SRC_ANT", r["id"], "E")] is not None:
            nm(wb, r["id"], f"SRC_ANT!$E${n}")
        if final[("SRC_ANT", r["id"], "F")] is not None:
            nm(wb, r["id"] + "_Lo", f"SRC_ANT!$F${n}")
        if final[("SRC_ANT", r["id"], "G")] is not None:
            nm(wb, r["id"] + "_Hi", f"SRC_ANT!$G${n}")
        if vals[9] is not None:
            nm(wb, r["id"] + "_Date", f"SRC_ANT!$J${n}")
            ws[f"J{n}"].number_format = "yyyy-mm-dd"
    for col, w in zip(["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "L", "M", "Q", "T", "V", "W"],
                      (13, 26, 5, 40, 9, 8, 8, 14, 14, 11, 50, 30, 30, 30, 16, 60)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "E5"


def sheet_tk(ctx):
    wb, snap, date = ctx.wb, ctx.snap, ctx.date
    ws = wb.create_sheet("TK_Link")
    title(ws, "TK_Link — Tokenomics 快照（第 0 層連結；不用 Excel 外部連結）",
          "每列一個 Tokenomics 名稱。值（藍字）由 builder 從 Tokenomics master 的 model/CURRENT 讀出寫入；本模型公式只引用 TK_ 具名範圍。前 63 名與 OpenAI v0.6 相同，其後 5 名為 A2 新增（任務別每次嘗試 token）；"
          "其後為 NonNV 表（Tokenomics 未設具名範圍，以列標籤＋欄標題讀表；狀態「讀表（非具名）」，列入 Tokenomics 缺口：建議 Tokenomics 把 NonNV 比例列進 Interface）。",
          "Interface 向量名稱為 15 欄＝5 世代（Hopper、GB200、GB300、VR200、Rubin Ultra）× 3 成本情境（低、基準、高）；NonNV 比例皆相對 VR200（同層級、同 SLO）。")
    meta = [("Tokenomics 檔案", snap["file"], "TK_File"), ("Tokenomics 版本", snap["version"], "TK_Version"),
            ("Tokenomics 提交 SHA（master）", snap["sha"], "TK_Commit"), ("讀取日期", date, "TK_ReadDate")]
    for i, (lab, v, name) in enumerate(meta):
        put(ws, f"A{3 + i}", lab, F_BOLD)
        put(ws, f"B{3 + i}", v, F_CALC)
        nm(wb, name, f"TK_Link!$B${3 + i}")
    L = openpyxl.utils.get_column_letter
    put(ws, "I7", "世代", F_BOLD)
    put(ws, "I8", "成本情境", F_BOLD)
    for k in range(15):
        put(ws, f"{L(TK_VCOL0 + k)}7", snap["hdr_gen"][k], F_CALC)
        put(ws, f"{L(TK_VCOL0 + k)}8", snap["hdr_cost"][k], F_CALC)
    nm(wb, "TK_HdrGen", "TK_Link!$J$7:$X$7")
    nm(wb, "TK_HdrCost", "TK_Link!$J$8:$X$8")
    header(ws, 9, ["Tokenomics 名稱", "本模型名稱", "說明", "單位", "欄數", "狀態", "Tokenomics 版本", "提交 SHA", "讀取日期"] + [f"值{k + 1}" for k in range(15)])
    r = TK_FIRST
    for row in snap["rows"] + snap.get("nonnv", []):
        n = len(row["values"])
        our = row.get("our", "TK_" + row["name"])
        vals = [row["name"], our, row["label"], row["unit"], n, row.get("status", "OK"), snap["version"], snap["sha"][:7], date]
        for c, v in enumerate(vals, start=1):
            put(ws, f"{L(c)}{r}", v, F_CALC)
        for k, v in enumerate(row["values"]):
            put(ws, f"{L(TK_VCOL0 + k)}{r}", v, F_IN)
        nm(wb, our, f"TK_Link!$J${r}" + (f":${L(TK_VCOL0 + n - 1)}${r}" if n > 1 else ""))
        r += 1
    for row in snap["pending"]:
        vals = [row["name"], "（Tokenomics 現行版無此名稱；無具名範圍）", "Tokenomics 缺口：見報告", "", 0, row["status"], snap["version"], snap["sha"][:7], date]
        for c, v in enumerate(vals, start=1):
            put(ws, f"{L(c)}{r}", v, F_CALC)
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 50
    ws.freeze_panes = "D10"


def sheet_oai(ctx):
    wb, oa = ctx.wb, ctx.oai
    ws = wb.create_sheet("OAI_Link")
    title(ws, "OAI_Link — OpenAI v0.6 命題輸出快照（規格 D20；只作並排對照）",
          "值由 builder 從 OpenAI v0.6 活頁簿的快取值讀出寫入（藍字）。本頁只被 Checks（與 A5 的 HTML）引用，任何計算頁不得引用（Checks 列 builder 掃描結果；tests 每次重驗）。",
          "OpenAI 的 2025 為其實際校準年；兩家共用同一 Tokenomics 快照（v5.26），故可直接並排。")
    meta = [("OpenAI 檔案", oa["file"], "OAI_File"), ("檔案 SHA-256", oa["sha256"], "OAI_SHA256"), ("git 提交（建置時）", oa["commit"], "OAI_Commit"),
            ("讀取日期", ctx.date, "OAI_ReadDate")]
    for i, (lab, v, name) in enumerate(meta):
        put(ws, f"A{4 + i}", lab, F_BOLD)
        put(ws, f"B{4 + i}", v, F_CALC)
        nm(wb, name, f"OAI_Link!$B${4 + i}")
    header(ws, 9, ["OpenAI 名稱", "本模型名稱", "說明", "2025", "2026", "2027", "2028", "2029", "2030", "單位", "OpenAI 位置"])
    for i, row in enumerate(oa["rows"]):
        n = 10 + i
        for c, v in enumerate([row["name"], "OAI_" + row["name"], row["label"]], start=1):
            ws.cell(row=n, column=c, value=v).font = F_CALC
        for k, v in enumerate(row["values"]):
            put(ws, f"{'DEFGHI'[k]}{n}", v, F_IN)
        put(ws, f"J{n}", row["unit"], F_CALC)
        put(ws, f"K{n}", row["ref"], F_NOTE)
        nm(wb, "OAI_" + row["name"], f"OAI_Link!$D${n}:$I${n}")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 34


def sheet_inputs(ctx, final):
    D, wb = ctx.D, ctx.wb
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs — 假設輸入（藍字）：值、低、高、標記、依據、區間理由",
          "長表格式：一列一個（參數 × 索引）。Analogy／Assumed 一律附區間（M 欄檢查）；Decision 為定義常數或開關（低＝高＝值）。來源 data/anthropic_inputs.yaml；ID 由 builder/id_registry.json 固定。",
          "引用 OpenAI v0.6 同類參數作 Analogy 者，依據欄寫明「OpenAI v0.6 INP_nnn」與差異。")
    header(ws, 4, INP_COLS)
    for i, r in enumerate(D.inp_rows):
        n = INP_FIRST + i
        vals = [r["id"], r["name"], s(r.get("index")) or "—", r["unit"], None, None, None, r["tag"], r.get("basis"), r.get("range_why"), r.get("note")]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=n, column=c, value=v).font = F_CALC
        for col in "EFG":
            v = final[("Inputs", r["id"], col)]
            put(ws, f"{col}{n}", v, F_IN)
        put(ws, f"L{n}", f"=IF(AND(ISNUMBER($E{n}),ISNUMBER($F{n}),ISNUMBER($G{n})),IF(AND($F{n}<=$E{n},$E{n}<=$G{n}),0,1),0)")
        put(ws, f"M{n}", f'=IF(AND(OR($H{n}="Analogy",$H{n}="Assumed"),OR(NOT(ISNUMBER($F{n})),NOT(ISNUMBER($G{n})))),1,0)')
        put(ws, f"N{n}", r["key"], F_NOTE)
        nm(wb, r["id"], f"Inputs!$E${n}")
    for col, w in zip("ABCDEFGHIJKN", (10, 44, 12, 12, 10, 9, 9, 13, 60, 40, 40, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "E5"


def base_checks(ctx, n_src, n_inp):
    snap = ctx.snap
    tbl = [r["our"] for r in snap.get("nonnv", [])]
    nnv = [x for x in tbl if x.startswith("TK_NNV_")]
    fams = sorted({x.split("_")[2] for x in nnv})
    order = "+".join(f"(TK_NNV_{f}_{a}Lo>TK_NNV_{f}_{a})+(TK_NNV_{f}_{a}>TK_NNV_{f}_{a}Hi)" for f in fams for a in ("Out", "Hold"))
    rows = [
        ("src_rows", "SRC_ANT 列數", "=SUMPRODUCT(--(LEN(SRC_ANT!$A$5:$A$600)>0))", n_src, "eq", "A1 376 列＋A2 新增 3 列（SRC_ANT_377–379）"),
        ("src_missing", "SRC_ANT 缺值列數", "=SUM(SRC_ANT!$X$5:$X$600)", 0, "eq", "值、低、高皆空且非定性／文字／日期列"),
        ("src_nosource", "SRC_ANT 缺出處或網址列數", "=SUM(SRC_ANT!$Y$5:$Y$600)", 0, "eq", ""),
        ("src_order", "SRC_ANT 區間順序異常列數", "=SUM(SRC_ANT!$Z$5:$Z$600)", 0, "eq", "低 ≤ 值 ≤ 高"),
        ("src_derived", "SRC_ANT 標記 Derived 但未處理（仍為觀測）列數", "=SUM(SRC_ANT!$AA$5:$AA$600)", 0, "eq", "yaml 檔頭：Derived 列改公式或移入 Inputs"),
        ("inp_rows", "Inputs 列數", "=SUMPRODUCT(--(LEN(Inputs!$A$5:$A$600)>0))", n_inp, "eq", ""),
        ("inp_order", "Inputs 區間順序異常列數", "=SUM(Inputs!$L$5:$L$600)", 0, "eq", "低 ≤ 值 ≤ 高（有三值者）"),
        ("inp_norange", "Inputs Analogy／Assumed 缺區間列數", "=SUM(Inputs!$M$5:$M$600)", 0, "eq", "共同規則第 4 節：Analogy／Assumed 一律給區間"),
        ("tk_ok", "TK_Link 具名範圍已取值名稱數", '=COUNTIF(TK_Link!$F$10:$F$200,"OK")', len(snap["rows"]), "eq", "OpenAI v0.6 同一組 63 名＋A2 新增 5 名（IF_HdrTask、IF_TaskLen、IF_TaskTok*）"),
        ("tk_nnv", "TK_Link 讀表（非具名）格數：NonNV 18＋PUE 3", f'=COUNTIF(TK_Link!$F$10:$F$200,"{NONNV_STATUS}")', len(tbl), "eq",
         "TPU v7、Trainium3、AMD MI455X × 6 值（規格 D6、D6 r1）；A3 加 Tokenomics Inputs 頁 PUE（低／基準／高）"),
        ("tk_pending", "TK_Link 待 Tokenomics 提供名稱數", f'=COUNTIF(TK_Link!$F$10:$F$200,"{PENDING_NOTE}")', len(snap["pending"]), "eq", ""),
        ("tk_err", "TK_Link 錯誤值格數", "=SUMPRODUCT(--ISERROR(TK_Link!$J$10:$X$200))", 0, "eq", ""),
        ("tk_nnv_order", "TK NonNV 產出比、持有比 低 ≤ 基準 ≤ 高 違反數", "=" + order if order else "=0", 0, "eq", "Tokenomics 快照被改寫時轉 ERR"),
        ("tk_pue_order", "TK PUE 低 ≤ 基準 ≤ 高 違反數", "=(TK_PUE_Lo>TK_PUE)+(TK_PUE>TK_PUE_Hi)", 0, "eq", "A3：PUE 取自 Tokenomics Inputs 頁（非具名）"),
        ("tk_version", "TK 快照：Tokenomics 版本", "=TK_Version", None, "info", "check_tk_snapshot 結果見報告"),
        ("tk_commit", "TK 快照：Tokenomics master 提交", "=TK_Commit", None, "info", ""),
        ("oai_isolated", "OAI_Link 被計算頁引用的公式格數（builder 掃描）", ctx.scan["oai_refs"], 0, "eq",
         "規格 D20：只被 Checks 引用；本列為 builder 建置時掃描結果，tests/parity/test_builder.py 每次重驗"),
        ("oai_sha", "OAI_Link：OpenAI 活頁簿 SHA-256", "=OAI_SHA256", None, "info", ""),
        ("e6", "E6：計算頁公式含常數的格數（builder 掃描）", ctx.scan["e6"], 0, "eq",
         "Demand、Revenue（A3 起含 Compute、Cost…）；本列為建置時掃描結果，tests 每次重驗"),
    ]
    return rows


def sheet_checks(ctx, rows):
    wb, reg = ctx.wb, ctx.D.reg
    ws = wb.create_sheet("Checks")
    title(ws, f"Checks — {STAGE} 檢查（公式；結果 ERR 的格數＝CHK_Errors，必須為 0；WARN 只提示）",
          "OK／ERR：比對期望值；WARN：超過門檻只提示（門檻在 Inputs）；INFO：只列示。編號（C##）由 builder/id_registry.json 固定：新檢查取下一個號碼，退役號碼不重用。")
    header(ws, 4, ["編號", "檢查項", "值（公式）", "期望", "結果", "說明"])
    known = reg["CHK"]
    nxt = max([int(v[1:]) for v in known.values()] + [int(x[1:]) for x in reg.get("retired_CHK", [])] + [0]) + 1
    for r_ in rows:
        if r_[0] not in known:
            known[r_[0]] = f"C{nxt:02d}"
            nxt += 1
    keys = {r_[0] for r_ in rows}
    for k in [k for k in known if k not in keys]:
        reg.setdefault("retired_CHK", []).append(known.pop(k))
    rows = sorted(rows, key=lambda r_: int(known[r_[0]][1:]))
    for i, (key, lab, f, exp, kind, note) in enumerate(rows):
        n = 5 + i
        put(ws, f"A{n}", known[key], F_CALC)
        put(ws, f"B{n}", lab, F_CALC)
        put(ws, f"C{n}", f, F_CALC)
        if kind == "warn":
            put(ws, f"D{n}", f"={exp}", F_CALC)
            put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(ABS(C{n})>D{n},"WARN","OK"),"ERR")')
        else:
            if exp is not None:
                put(ws, f"D{n}", exp, F_CALC)
            if kind == "eq":
                put(ws, f"E{n}", f'=IF(C{n}=D{n},"OK","ERR")')
            elif kind == "tol":
                put(ws, f"E{n}", f'=IF(ISNUMBER(C{n}),IF(ABS(C{n}-D{n})<0.000000001,"OK","ERR"),"ERR")')
            else:
                put(ws, f"E{n}", "INFO")
        put(ws, f"F{n}", note, F_NOTE)
    last = 4 + len(rows)
    put(ws, f"A{last + 2}", "CHK_Errors", F_BOLD)
    put(ws, f"C{last + 2}", f'=COUNTIF($E$5:$E${last},"ERR")')
    nm(wb, "CHK_Errors", f"Checks!$C${last + 2}")
    put(ws, f"A{last + 3}", "CHK_Warnings", F_BOLD)
    put(ws, f"C{last + 3}", f'=COUNTIF($E$5:$E${last},"WARN")')
    nm(wb, "CHK_Warnings", f"Checks!$C${last + 3}")
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["F"].width = 80
    return {known[r_[0]]: r_[0] for r_ in rows}


def scan(wb):
    """E6 與 OAI_Link 隔離：建置時掃描（tests 另行重驗）。"""
    e6, oai = [], []
    for ws in wb.worksheets:
        if ws.title in CALC_SHEETS:
            e6 += e6_violations(ws)
        if ws.title in ("OAI_Link", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and ("OAI_" in c.value or "OAI_Link!" in c.value):
                    oai.append(f"{ws.title}!{c.coordinate}")
    rvs = []                                          # D19：Reverse 只被 Checks 引用（不回饋基準）
    for ws in wb.worksheets:
        if ws.title in ("Reverse", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and ("Reverse!" in c.value or "RVS_" in c.value):
                    rvs.append(f"{ws.title}!{c.coordinate}")
    return dict(e6=len(e6), e6_cells=e6, oai_refs=len(oai), oai_cells=oai, rvs_refs=len(rvs), rvs_cells=rvs)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tk-dir", required=True, type=Path)
    ap.add_argument("--oai-xlsx", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--base", type=Path)
    ap.add_argument("--date", default=dt.date.today().isoformat())
    ap.add_argument("--no-lo", action="store_true", help="略過 LibreOffice 重算（輸出無快取值；僅供除錯）")
    a = ap.parse_args()
    D = Data()
    snap = read_snapshot(a.tk_dir)
    snap["nonnv"] = read_nonnv(a.tk_dir)
    oai = read_oai(a.oai_xlsx)
    outdir = ROOT / "build_out"
    outdir.mkdir(parents=True, exist_ok=True)
    a.out.parent.mkdir(parents=True, exist_ok=True)

    defaults = {}
    for r in D.src_rows:
        f = DERIVED.get(r["key"])
        defaults[("SRC_ANT", r["id"], "E")] = D.fill(f) if f else blank(r["value"])
        defaults[("SRC_ANT", r["id"], "F")] = None if f else blank(r["lo"])
        defaults[("SRC_ANT", r["id"], "G")] = None if f else blank(r["hi"])
    for r in D.inp_rows:
        for col, k in zip("EFG", ("value", "lo", "hi")):
            defaults[("Inputs", r["id"], col)] = blank(r.get(k))
    if a.base:
        final, matched, kept, unmatched = preserve.merge(defaults, a.base, outdir / "restore_log.txt")
    else:
        final, matched, kept, unmatched = dict(defaults), 0, 0, []
        (outdir / "restore_log.txt").write_text("matched: 0; Excel value kept over code default: 0; unmatched: 0\n（首次建置，無底稿）\n", encoding="utf-8")
    print(f"restore: matched {matched}, Excel kept over code {kept}, unmatched {len(unmatched)}")

    wb = openpyxl.Workbook()
    ctx = Ctx(wb, D, snap, oai, a.date)
    summary = dict(src=len(D.src_rows), inp=len(D.inp_rows), tk_ok=len(snap["rows"]), tk_nnv=len(snap["nonnv"]), tk_pending=len(snap["pending"]))
    sheet_readme(ctx, summary)
    sheet_src(ctx, final)
    sheet_tk(ctx)
    sheet_oai(ctx)
    sheet_inputs(ctx, final)
    mods = [importlib.import_module(m) for m in STAGE_MODULES if (HERE / f"{m}.py").exists()]
    for m in mods:
        m.build(ctx)
    ctx.scan = scan(wb)
    rows = base_checks(ctx, len(D.src_rows), len(D.inp_rows))
    for m in mods:
        rows += m.checks(ctx)
    ids = sheet_checks(ctx, rows)
    preserve.write_defaults(wb, final)
    D.save_registry()
    raw = Path(tempfile.mkdtemp()) / a.out.name
    wb.save(raw)
    if a.no_lo:
        raw.replace(a.out)
    else:
        lo_recalc(raw, Path(tempfile.mkdtemp())).replace(a.out)
    json.dump(dict(summary, stage=STAGE, modules=[m.__name__ for m in mods], tk_file=snap["file"], tk_sha=snap["sha"], oai_file=oai["file"],
                   oai_sha256=oai["sha256"], oai_commit=oai["commit"], scan={k: v for k, v in ctx.scan.items()}, checks=ids),
              open(outdir / "build_summary.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("saved", a.out, summary, "E6", ctx.scan["e6"], "OAI refs", ctx.scan["oai_refs"])


if __name__ == "__main__":
    main()
