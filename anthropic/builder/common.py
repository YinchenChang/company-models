# 共用樣式與函式（沿用 Tokenomics builder/common.py 的慣例：藍字＝輸入、黑字＝公式、綠字＝跨頁連結）
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L  # noqa: F401

BLUE, BLACK, GREEN = "FF0000FF", "FF000000", "FF008000"
F_IN = Font(name="Arial", size=10, color=BLUE)
F_CALC = Font(name="Arial", size=10, color=BLACK)
F_LINK = Font(name="Arial", size=10, color=GREEN)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_TITLE = Font(name="Arial", size=13, bold=True)
F_NOTE = Font(name="Arial", size=9, color="FF595959")
FILL_SEC = PatternFill("solid", fgColor="FFD9E1F2")
FILL_HDR = PatternFill("solid", fgColor="FFEDEDED")
WRAP = Alignment(wrap_text=True, vertical="top")


def put(ws, ref, v, font=None, fmt=None, fill=None, wrap=False):
    c = ws[ref]
    c.value = v
    if font is None:
        if isinstance(v, str) and v.startswith("="):
            font = F_LINK if "!" in v else F_CALC
        else:
            font = F_CALC
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    return c


def title(ws, t1, t2, t3=None):
    put(ws, "A1", t1, F_TITLE)
    put(ws, "A2", t2, F_NOTE)
    if t3:
        put(ws, "A3", t3, F_NOTE)


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = F_BOLD
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def lo_recalc(xlsx: Path, outdir: Path) -> Path:
    """以 LibreOffice headless 重算並另存 xlsx（與 tests/parity/parity_lib.py 相同做法），回傳輸出路徑。"""
    outdir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    try:
        cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", str(outdir), str(xlsx)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / xlsx.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice 重算失敗：{r.stdout}\n{r.stderr}")
    return out


def nm(wb, name, ref):
    """新增活頁簿層級具名範圍。"""
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def e6_violations(ws, min_row=5):
    """E6：公式不得內含常數（恆等式 1−比例、1＋成長率、年數 +1 除外）。回傳違反的儲存格清單（與 tests/parity/test_builder.py 同一規則）。"""
    import re
    bad = []
    for row in ws.iter_rows(min_row=min_row):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                f = re.sub(r"'?[A-Za-z_]+'?!\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?", "", c.value)
                f = re.sub(r"\$?\b[A-Z]{1,3}\$?\d+\b", "", f)
                f = re.sub(r"\b[A-Za-z_][A-Za-z_0-9]*\b", "", f)
                f = re.sub(r'"[^"]*"', "", f)
                f = f.replace("(1-", "(").replace("+1)", ")").replace("(1+", "(")
                if re.search(r"\d", f):
                    bad.append(f"{ws.title}!{c.coordinate}")
    return bad
