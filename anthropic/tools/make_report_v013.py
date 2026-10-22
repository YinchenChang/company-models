#!/usr/bin/env python3
"""v0.1.3 A 段對照 Excel（CLAUDE.md 第 5 節）：①新增或變動的列 ②關鍵輸出前後（v0.1.2 → v0.1.3）③已套用的預設。

用法：python3 tools/make_report_v013.py --old model/archive/20261021_Anthropic_v0.1.2.xlsx --new model/20261022_Anthropic_v0.1.3.xlsx \
        --out docs/reports/20261022_v0.1.3-A_對照.xlsx
數值一律讀兩個活頁簿 LibreOffice 重算後的快取值；本程式不做模型計算。
"""
import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="FFEDEDED")
YEARS = ["2025", "2026", "2027", "2028", "2029", "2030"]

# ② 關鍵輸出（具名範圍, 顯示名稱）
KEY_OUT = [
    ("COST_PropGap_VR", "每 VR 等值 GW 差額（含股權報酬，$B/GW/年）"), ("COST_PropRev_VR", "每 VR 等值 GW 營收淨額（$B/GW/年）"),
    ("COST_PropFull_VR", "每 VR 等值 GW 全成本（$B/GW/年）"), ("COST_Coverage", "覆蓋率（營收淨額 ÷ 全成本）"),
    ("CMP_Supply_VReq", "供給 VR 等值 GW"), ("CMP_SupplyGW", "供給 GW"), ("CMP_Supply_Nscale", "Nscale 供給 GW"), ("CMP_Supply_Lambda", "Lambda 供給 GW"),
    ("COST_Pay_Nscale", "Nscale 合約實付（$B）"), ("COST_Pay_Lambda", "Lambda 合約實付（$B）"), ("COST_ContractPay", "合約實付合計（$B）"),
    ("COST_Compute", "算力成本（現金，$B）"), ("REV_NetCapped", "營收淨額（截頂後，$B）"), ("FND_FCF", "營運自由現金流（$B）"),
    ("FND_CashEnd", "年底現金（基準，$B）"), ("FND_ExtNeedCum", "累計外部資金需求（基準，$B）"),
]
# 本版新增的單一輸出（只有新版有）
NEW_SCALAR = [
    ("FND_FCFFirstPos", "營運自由現金流首次轉正年"),
    ("FND_L1_GapYear", "第 1 層（營運自給）最低現金缺口首次出現年"), ("FND_L1_ExtTotal", "第 1 層累計外部資金需求 2026–2030（$B）"), ("FND_L1_Trough", "第 1 層現金谷底（$B）"),
    ("FND_L2_GapYear", "第 2 層（加已到位股權＝基準）缺口年"), ("FND_L2_ExtTotal", "第 2 層累計外部資金需求（$B）"), ("FND_L2_Trough", "第 2 層現金谷底（$B）"),
    ("FND_L3_GapYear", "第 3 層（再加條件式）缺口年"), ("FND_L3_ExtTotal", "第 3 層累計外部資金需求（$B）"), ("FND_L3_Trough", "第 3 層現金谷底（$B）"),
]
EXP = [("None", "無需求（變現 0）"), ("Half", "部分變現（0.5）"), ("Full", "完全變現（1）")]
DEFAULTS = [
    ("U2 新法：付款路徑", "chat 端更正：付款路徑不變＝『最高 86』按 MW 比例（460:350）分攤、÷ Σ 爬坡權重 × 當年權重（v0.1.1／v0.1.2 的路徑）；只把供給 GW 改回揭露容量", "按兩約公告金額比例（45:35）分攤（工作單文字）", "由比例推得：Nscale 年付款約 −1.0%（8.88 → 約 8.80）、Lambda 約 +1.3%；對差額的影響未另行重算。未採用，因工作單同一句要求『v0.1.2 付款路徑不變』"),
    ("U2 新法：Lambda 供給期間", "起始年（Inputs）至 WSJ 轉述的結束年 2034", "沿用 v0.1 的期間 6 年（2027–2032）", "2030 以前相同"),
    ("U2 新法開關", "Inputs u2_method（INP_187）＝1；0 且 u2_switch＝1 重現 v0.1.1／v0.1.2", "無", "2030 差額 11.14 → 11.60；成本不變"),
    ("資金第 1 層是否計 2025 年底現金", "不計（Inputs l1_incl_cash＝0；該現金是過去融資所得）", "計入（＝1）", "計入則第 1 層起算現金多 20.28，累計外部資金需求最多減少同額"),
    ("資金第 3 層內容", "IPO＋Amazon＋Google＋NVIDIA＋AMD 全數計入（上緣）", "NVIDIA 視為 IPO 基石不另計（既有 S3 的處理）", "兩種在基準下都無缺口，結論不變"),
    ("持續擴張：追加金額", "每個追加年（2029、2030）＝當年基準合約付款 × 30%（Assumed，0–60%），為該批滿載年付款，依爬坡年數爬坡", "以新增 GW 目標反推", "2030 追加實付 43.0（基準合約付款 92.3 的 47%）"),
    ("持續擴張：變現", "0／0.5／1 三檔；新增 GW 營收＝變現比例 × 追加 VR 等值 GW × 當年基準每 VR 等值 GW 營收淨額，不受容量上限約束", "經容量上限與價格彈性重算", "完全變現為上緣；基準模型不受影響"),
    ("TK v19.0 標記", "快照改 v19.0；103 個值與 v18.0 相同，模型數值不變", "維持 v18.0", "無"),
    ("OpenAI 對照快照", "維持 v0.6（main 現行為 v0.6.1，不在本版範圍）", "改讀 v0.6.1", "只影響並排列"),
]


def named(wb, n):
    sh, rg = wb.defined_names[n].attr_text.split("!")
    c = wb[sh.strip("'")][rg.replace("$", "")]
    if not isinstance(c, tuple):
        c = ((c,),)
    return [x.value for r in c for x in r]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--old", required=True, type=Path)
    ap.add_argument("--new", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    a = ap.parse_args()
    old, new = openpyxl.load_workbook(a.old, data_only=True), openpyxl.load_workbook(a.new, data_only=True)
    out = openpyxl.Workbook()

    # ① 新增或變動的列
    ws = out.active
    ws.title = "①新增或變動的列"
    ws.append(["頁", "編號／ID", "列名", "單位", "2025", "2026", "2027", "2028", "2029", "2030", "標記／說明"])
    for sh in ("Compute", "Cost", "Funding"):
        o, n = old[sh], new[sh]
        last = max((r for r in range(7, o.max_row + 1) if o.cell(r, 1).value), default=0)
        for r in range(last + 1, n.max_row + 1):
            if n.cell(r, 2).value and n.cell(r, 1).value and not str(n.cell(r, 1).value).startswith("（"):
                ws.append([sh, n.cell(r, 1).value, n.cell(r, 2).value, n.cell(r, 3).value] + [n.cell(r, c).value for c in range(4, 10)] + [n.cell(r, 11).value])
    ws.append([])
    ws.append(["Inputs（新增）", "ID", "參數", "單位", "值", "低", "高", "標記", "依據"])
    oi = {r[0].value for r in old["Inputs"].iter_rows(min_row=5) if r[0].value}
    for r in new["Inputs"].iter_rows(min_row=5):
        if r[0].value and r[0].value not in oi:
            ws.append(["Inputs", r[0].value, r[1].value, r[3].value, r[4].value, r[5].value, r[6].value, r[7].value, r[8].value])
    ws.append([])
    ws.append(["Checks（新增）", "ID", "檢查項", "值", "期望", "狀態"])
    oc = {r[0].value for r in old["Checks"].iter_rows(min_row=5) if r[0].value}
    for r in new["Checks"].iter_rows(min_row=5):
        if r[0].value and r[0].value not in oc:
            ws.append(["Checks", r[0].value, r[1].value, r[2].value, r[3].value, r[4].value])

    # ② 關鍵輸出前後
    w2 = out.create_sheet("②關鍵輸出前後")
    w2.append(["輸出", "版本"] + YEARS + ["說明"])
    for n_, lab in KEY_OUT:
        ov, nv = named(old, n_), named(new, n_)
        w2.append([lab, "v0.1.2"] + ov)
        w2.append([lab, "v0.1.3"] + nv)
        w2.append([lab, "差（新 − 舊）"] + [(b - a_) if isinstance(a_, (int, float)) and isinstance(b, (int, float)) else None for a_, b in zip(ov, nv)])
    w2.append([])
    w2.append(["資金三層與首次轉正年（v0.1.3 新增）", "", "值"])
    for n_, lab in NEW_SCALAR:
        w2.append([lab, "v0.1.3", named(new, n_)[0]])
    w2.append([])
    w2.append(["持續擴張情境（v0.1.3 新增；基準不變）", "指標"] + YEARS)
    for k, lab in EXP:
        w2.append([f"擴張・{lab}", "每 VR 等值 GW 差額"] + named(new, f"COST_PropGap_VR_Exp{k}"))
        w2.append([f"擴張・{lab}", "覆蓋率"] + named(new, f"COST_Coverage_Exp{k}"))
        w2.append([f"擴張・{lab}", "營運自由現金流 $B"] + named(new, f"FND_FCF_Exp{k}"))
        w2.append([f"擴張・{lab}", "追加實付 $B"] + named(new, "COST_ExpPay"))
        w2.append([f"擴張・{lab}", "首次轉正年 / 第1層外部資金 / 第2層 / 第3層 / 第2層谷底"] + [named(new, f"FND_FCFFirstPos_Exp{k}")[0], named(new, f"FND_L1_ExtTotal_Exp{k}")[0],
                                                                   named(new, f"FND_L2_ExtTotal_Exp{k}")[0], named(new, f"FND_L3_ExtTotal_Exp{k}")[0], named(new, f"FND_L2_Trough_Exp{k}")[0]])

    # ③ 已套用的預設
    w3 = out.create_sheet("③已套用的預設")
    w3.append(["問題", "採用的預設", "替代", "對結果的影響"])
    for d in DEFAULTS:
        w3.append(list(d))

    for w in out.worksheets:
        for c in w[1]:
            c.font, c.fill = BOLD, HDR
        w.column_dimensions["A"].width = 44
        w.column_dimensions["B"].width = 40
        w.column_dimensions["C"].width = 40
        for col in "DEFGHIJ":
            w.column_dimensions[col].width = 16
        for row in w.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(c.value, float):
                    c.number_format = "#,##0.000"
    a.out.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
