#!/usr/bin/env python3
"""v0.1.6 對照 Excel（CLAUDE.md 第 5 節、第 8 節）：
①新增或變動的列（SRC_ANT_410–464、Inputs INP_195–196、Revenue h1act、Cost 第二十四節、Checks f1_*；頁、列名、ID、值、來源／標記）
②關鍵輸出 FY2025–2030：v0.1.5 vs v0.1.6（逐年）＋B 情境（η＝2026H1）
③已套用的預設
④S-1 替換對照（CLAUDE.md 第 8 節第 4 步：舊列 ID／舊值／新列 ID／新值／差異比例／是否進基準）
⑤逐格比較（全頁；分類與範圍外差異）
用法：python3 tools/make_report_v016.py --new model/20261031_Anthropic_v0.1.6.xlsx --old model/archive/20261024_Anthropic_v0.1.5.xlsx \
        --out docs/reports/20261031_v0.1.6_對照.xlsx [--json out.json]
數值一律讀 LibreOffice 重算後的快取值；本程式不做模型計算（差、比例只用於對照表的顯示）。
"""
import argparse
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

KEY_OUT = [("REV_GrossCapped", "營收總額（$B）"), ("REV_NetCapped", "營收淨額（$B）"), ("COST_PropRev_VR", "每 VR 等值 GW 營收淨額（$B/GW/年）"),
           ("COST_PropFull_VR", "每 VR 等值 GW 全成本"), ("COST_PropGap_VR", "每 VR 等值 GW 差額（$B/GW/年）"), ("COST_Coverage", "覆蓋率"),
           ("CMP_Supply_VReq", "供給 VR 等值 GW"), ("CMP_Eta", "η"), ("CMP_InfGW_Eff", "有效推論 GW"), ("CMP_RDGW", "研發 GW"), ("COST_Compute", "算力成本（$B）"),
           ("FND_FCF", "營運自由現金流（$B）"), ("FND_CashEnd", "年底現金（$B）"), ("FND_ExtNeedCum", "累計外部資金需求（$B）")]
SCEN = [("CMP_CapFactor_EtaH1", "B 情境：容量上限係數"), ("COST_PropRev_VR_EtaH1", "B 情境：每 VR 等值 GW 營收淨額"),
        ("COST_PropGap_VR_EtaH1", "B 情境：每 VR 等值 GW 差額"), ("COST_Coverage_EtaH1", "B 情境：覆蓋率")]
SCALARS = [("FND_CashNoExtTrough", "只靠已到位融資的現金谷底"), ("FND_L1_ExtTotal", "第 1 層累計外部資金需求"), ("CMP_Eta2025", "η（2025）"),
           ("COST_Eta2026H1", "η_2026H1（W3 c）"), ("COST_Eta2025COR", "η（2025）以 S-1 營收成本換算（對照）"), ("COST_S1fGapPayH1", "W3(a) 合約實付 ÷ S-1 − 1"),
           ("COST_S1fGapCcH1", "W3(a) 算力成本（含自建）÷ S-1 − 1"), ("COST_S1fGapCOR2025", "W3(b) S-1 營收成本 ÷ 模型推論 − 1"),
           ("COST_S1fCommitGap", "W3(d) 模型全部新合約 ÷ S-1 三段 − 1"), ("COST_S1fNamedGap", "W3(d) 模型七家 ÷ S-1 七家 − 1"), ("COST_PartnerGapMaxS1f", "W3(d) 逐家最大 |差距|")]
# CLAUDE.md 第 8 節第 4 步：舊列（草擬版）→ 新列（正式 S-1）
REPL = [("SRC_ANT_380", "SRC_ANT_414", "是（W2：Revenue h1act → REV_Target 2026）"), ("SRC_ANT_381", "SRC_ANT_413", "否"), ("SRC_ANT_382", "SRC_ANT_420", "否"),
        ("SRC_ANT_383", "SRC_ANT_421", "否"), ("SRC_ANT_004", "SRC_ANT_419", "否（U1 開關 0 路徑與 a3 對照列仍用舊列）"), ("SRC_ANT_005", "SRC_ANT_420", "否（同上）"),
        ("SRC_ANT_006", "SRC_ANT_416", "否"), ("SRC_ANT_001", "SRC_ANT_412", "否（工作單：2025 校準不改）"), ("SRC_ANT_003", "SRC_ANT_411", "否"),
        ("SRC_ANT_252", "SRC_ANT_424", "否（η 校準與 2025 算力不改）"), ("SRC_ANT_394", "SRC_ANT_426", "否（U4 對照不改；W3 a 另列）"), ("SRC_ANT_395", "SRC_ANT_425", "否"),
        ("SRC_ANT_278", "SRC_ANT_433", "否（正負號口徑相反：舊列以虧損正值登錄）"), ("SRC_ANT_397", "SRC_ANT_450", "否（U3 對照不改；W3 d 另列）"),
        ("SRC_ANT_398", "SRC_ANT_451", "否"), ("SRC_ANT_399", "SRC_ANT_452", "否"), ("SRC_ANT_229", "SRC_ANT_453", "否"), ("SRC_ANT_401", "SRC_ANT_454", "否"),
        ("SRC_ANT_400", "SRC_ANT_455", "否（舊列為 Nscale＋Lambda 合計）"), ("SRC_ANT_232", "SRC_ANT_457", "否"), ("SRC_ANT_405", "SRC_ANT_461", "否（口徑不同：兩家 vs 單一最大）")]
DEFAULTS = [
    ("η 改用 2026H1", "只做情境（Cost 第二十四節 B 情境）；基準 η 不改", "基準 η 改用 η_2026H1", "0（η_2026H1 0.604 > 基準 0.525，容量不受限，營收、差額、覆蓋率不變；研發 GW 2030 增加）"),
    ("期中未查核數字的標記", "Verified，備註「期中數字，未經查核」", "Interested-party", "無數值影響"),
    ("S-1 留白欄位（Q3 營收、承諾總額、付款時程、短期承諾）", "記為「S-1 未揭露」（SRC 值空或 not_found），不填推論值", "以其他數字推算", "無"),
    ("sec.gov 無法直接讀到季度、Note 8、承諾章節", "改用二手轉載（Kepler、sec-api.io、Tunguz、AlphaSense、xenospectrum、Daloopa、thevccorner），標 Interested-party，報告列為缺口", "等 chat 端核對後再登錄", "只影響對照列；W2 基準所用 16.286 為一手直讀"),
    ("2026 Q2 營收兩個轉載值（11.555／11.560）", "兩者並列（SRC_ANT_420、421）；Q1＋Q2 核對用 Kepler 11.555（與上半年 16.286 相加一致）", "只登工作單的 11.56", "無（不進基準）"),
    ("W2 實作方式", "Inputs s1f_switch（INP_195）＝1 讀正式 S-1；0＝草擬版（v0.1.5 逐格重現）", "直接改公式、不設開關", "無；可回歸"),
    ("W3(c) 上半年 token 換算 GW", "全年 token 換算 GW × 上半年營收占比（S-1 上半年 ÷ 模型 2026 總額）× 12 ÷ 6（年內 token 與營收同比例）", "不折上半年（全年 token ÷ 上半年支出）", "η_2026H1 0.604；替代 0.925（上緣）"),
    ("W3(c) 合約價", "既有雲端加權合約價（2025 η 同一值 9.36）", "2026 年加權合約價（新合約組合）", "與 2025 同法；替代會改變 η 水準"),
    ("W3(c) B 情境的 η 路徑", "2025＝校準值、2026 起＝η_2026H1（各年同值）", "2025 也改 η_2026H1", "2025、2026 容量固定 1，無差別"),
    ("W3(d) 對照口徑", "S-1 三段合計（6-30 未起租租賃＋6-30 固定或最低承諾＋6-30 後新簽）vs 模型全部新合約（Cost 第十一節）；逐家同第十二節（模型 ÷ S-1 − 1）", "含已起租租賃或 Note 16 的 417（起算日 2025-12-31，與三段重疊）", "只對照"),
    ("W3(d) 逐家差距的判定", "INFO（不新增 WARN；第十二節草擬版 C89 WARN 不變）", "超過 25% 立 WARN", "CHK_Warnings 維持 2"),
    ("2025 校準營收 4.6 vs 正式 4.585、2025 算力 7.33 vs 7.326", "不改（工作單：基準只改 W2）；列 Checks f1_rev25 與報告", "改用正式值", "若改：2025 營收 −0.3%、η 微動"),
    ("OpenAI 並排快照", "維持 OAI_Link v0.6（2026 每 VR 等值 GW 營收 23.7）；報告與 HTML 另以文字列出 OpenAI 現行 v0.6.4 的 21.89", "OAI_Link 改 v0.6.4", "只影響並排列"),
    ("Tokenomics 快照", "維持 v21.2（master 45ce1ca）；TK master 現為 v21.7，本工作單未要求同步", "同步 v21.7", "無（未做）"),
    ("README 新列位置", "「v0.1.6 更新」放在 README 說明列的後段（層級對應之前）", "放最前", "僅版面（README 其後 3 列下移）"),
]


def named(wb, name):
    sh, rg = wb.defined_names[name].attr_text.replace("$", "").split("!")
    c = wb[sh.strip("'")][rg]
    if not isinstance(c, tuple):
        c = ((c,),)
    return [x.value for r in c for x in r]


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def same(a, b):
    return a == b or (num(a) and num(b) and abs(a - b) <= 1e-9 * max(1, abs(a), abs(b)))


def src_index(wb):
    ws = wb["SRC_ANT"]
    return {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("SRC_ANT_")}


def main():
    ap = argparse.ArgumentParser()
    for k in ("new", "old", "out"):
        ap.add_argument("--" + k, required=True, type=Path)
    ap.add_argument("--json", type=Path)
    a = ap.parse_args()
    nv, ov = openpyxl.load_workbook(a.new, data_only=True), openpyxl.load_workbook(a.old, data_only=True)
    nf, of = openpyxl.load_workbook(a.new), openpyxl.load_workbook(a.old)
    out = openpyxl.Workbook()
    B = Font(bold=True)
    summary = {}

    # ① 新增或變動的列
    ws = out.active
    ws.title = "①新增與變動列"
    ws.append(["頁", "列名", "ID", "值（2025／2026 或單值）", "低", "高", "來源／標記", "備註"])
    s = nv["SRC_ANT"]
    idx = src_index(nv)
    for sid, r in idx.items():
        if int(sid[8:]) >= 410:
            ws.append(["SRC_ANT", s.cell(r, 4).value, sid, s.cell(r, 5).value, s.cell(r, 6).value, s.cell(r, 7).value,
                       f"{s.cell(r, 16).value}｜{s.cell(r, 12).value}", s.cell(r, 23).value])
    oidx = src_index(ov)
    for sid, r in idx.items():
        if sid in oidx and s.cell(r, 23).value != ov["SRC_ANT"].cell(oidx[sid], 23).value:
            ws.append(["SRC_ANT（備註）", s.cell(r, 4).value, sid, s.cell(r, 5).value, None, None, "值不改；備註加註對應正式列", s.cell(r, 23).value[-40:]])
    inp = nv["Inputs"]
    for r in range(5, inp.max_row + 1):
        iid = inp.cell(r, 1).value
        if isinstance(iid, str) and iid.startswith("INP_") and int(iid[4:]) >= 195:
            ws.append(["Inputs", inp.cell(r, 2).value, iid, inp.cell(r, 5).value, inp.cell(r, 6).value, inp.cell(r, 7).value, inp.cell(r, 8).value, inp.cell(r, 9).value])
    rv, rf = nv["Revenue"], nf["Revenue"]
    for r in range(1, rv.max_row + 1):
        if str(rv.cell(r, 2).value or "").startswith("2026 上半年實際營收"):
            ws.append(["Revenue", rv.cell(r, 2).value, rv.cell(r, 1).value, rv.cell(r, 5).value, None, None, "公式：" + str(rf.cell(r, 5).value), "W2"])
    cv = nv["Cost"]
    start = next(r for r in range(1, cv.max_row + 1) if str(cv.cell(r, 1).value or "").startswith("二十四、"))
    for r in range(start + 1, cv.max_row + 1):
        vals = [cv.cell(r, c).value for c in range(4, 10)]
        shown = [v for v in vals if v is not None]
        ws.append(["Cost", cv.cell(r, 2).value, cv.cell(r, 1).value, "／".join(f"{v:.4g}" if num(v) else str(v) for v in shown[:6]), None, None, "Derived（只對照）", cv.cell(r, 11).value])
    ck = nv["Checks"]
    for r in range(5, ck.max_row + 1):
        cid = ck.cell(r, 1).value
        if isinstance(cid, str) and cid.startswith("C") and cid[1:].isdigit() and int(cid[1:]) >= 157:
            ws.append(["Checks", ck.cell(r, 2).value, cid, ck.cell(r, 3).value, None, None, ck.cell(r, 5).value, ck.cell(r, 6).value])

    # ② 關鍵輸出
    w2 = out.create_sheet("②關鍵輸出")
    w2.append(["輸出", "版本", "2025", "2026", "2027", "2028", "2029", "2030"])
    for nm, lab in KEY_OUT:
        o, n = named(ov, nm), named(nv, nm)
        w2.append([lab, "v0.1.5"] + o)
        w2.append([lab, "v0.1.6"] + n)
        w2.append([lab, "差"] + [(y - x) if num(x) and num(y) else None for x, y in zip(o, n)])
    w2.append([])
    w2.append(["B 情境（η＝2026H1；只作情境）", "v0.1.6"])
    for nm, lab in SCEN:
        w2.append([lab, "v0.1.6"] + named(nv, nm))
    w2.append([])
    w2.append(["標量", "v0.1.5", "v0.1.6", "差"])
    for nm, lab in SCALARS:
        n = named(nv, nm)[0]
        o = named(ov, nm)[0] if nm in ov.defined_names else None
        w2.append([lab, o, n, (n - o) if num(o) and num(n) else None])
    summary["gap"] = dict(old=named(ov, "COST_PropGap_VR"), new=named(nv, "COST_PropGap_VR"))

    # ③ 已套用的預設
    w3 = out.create_sheet("③已套用的預設")
    w3.append(["問題", "採用的預設", "替代選項", "對結果的影響方向"])
    for row in DEFAULTS:
        w3.append(list(row))

    # ④ S-1 替換對照
    w4 = out.create_sheet("④S-1替換對照")
    w4.append(["舊列 ID（草擬版）", "舊列項目", "舊值", "新列 ID（正式 S-1）", "新列項目", "新值", "差異比例（新 ÷ 舊 − 1）", "是否進基準"])
    for old_id, new_id, base in REPL:
        ro, rn = idx[old_id], idx[new_id]
        vo = s.cell(ro, 5).value if s.cell(ro, 5).value is not None else (s.cell(ro, 7).value if s.cell(ro, 7).value is not None else s.cell(ro, 6).value)
        vn = s.cell(rn, 5).value
        w4.append([old_id, s.cell(ro, 4).value, vo, new_id, s.cell(rn, 4).value, vn, (vn / vo - 1) if num(vo) and num(vn) and vo else None, base])

    # ⑤ 逐格比較
    w5 = out.create_sheet("⑤逐格比較")
    w5.append(["頁", "格", "類別", "舊值", "新值"])
    cats, tot = {}, 0
    for sh in nv.sheetnames:
        x, y, fx, fy = ov[sh], nv[sh], of[sh], nf[sh]
        for r in range(1, max(x.max_row, y.max_row) + 1):
            for c in range(1, max(x.max_column, y.max_column) + 1):
                tot += 1
                v1, v2, f1, f2 = x.cell(r, c).value, y.cell(r, c).value, fx.cell(r, c).value, fy.cell(r, c).value
                if same(v1, v2) and f1 == f2:
                    continue
                if sh == "_Defaults":
                    cat = "_Defaults（builder 預設值紀錄；隨新列增加）"
                elif sh == "SRC_ANT":
                    cat = "SRC_ANT 新列（410–464）" if r >= 5 + 409 else "SRC_ANT 舊列備註（值不改）"
                elif sh == "Inputs":
                    cat = "Inputs 新列（INP_195–196）"
                elif sh == "Cost" and r >= start:
                    cat = "Cost 第二十四節（新節，只對照）"
                elif sh == "Checks":
                    cat = "Checks（新增 C157 起；值與說明）"
                elif isinstance(v1, str) and isinstance(v2, str) and not str(f2 or "").startswith("="):
                    cat = "說明文字／版本字串／日期"
                elif sh in ("Demand", "Revenue", "Compute", "Cost", "Bridge", "EtaDiag", "Funding", "Reverse"):
                    cat = "W2 連動（2026 上半年 16.3 → 16.286 的下游數值）"
                else:
                    cat = "範圍外"
                cats[cat] = cats.get(cat, 0) + 1
                w5.append([sh, y.cell(r, c).coordinate, cat, v1 if not isinstance(v1, str) else v1[:120], v2 if not isinstance(v2, str) else v2[:120]])
    w5.append([])
    w5.append(["比較格數（所有頁）", tot])
    for k, v in cats.items():
        w5.append([k, v])
    summary["cells_total"], summary["cats"] = tot, cats
    for w in out.worksheets:
        for c in w[1]:
            c.font = B
        w.column_dimensions["A"].width = 30
        w.column_dimensions["B"].width = 50
    out.save(a.out)
    if a.json:
        a.json.write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print("saved", a.out, json.dumps(summary["cats"], ensure_ascii=False), tot)


if __name__ == "__main__":
    main()
