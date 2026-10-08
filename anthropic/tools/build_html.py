#!/usr/bin/env python3
"""HTML 一頁摘要產生器（A5；改寫自 OpenAI v0.6 的 tools/build_html.py）：
LibreOffice 重算 → 以具名範圍或「頁!列 ID」取值 → 單一靜態 HTML。

- 計算一律由 Excel 執行：基準與敏感度情境都是「改寫 Inputs（或 TK NonNV 快照格）→ LibreOffice 重算」後讀值；
  本程式只做取值、格式化與排版（含 SVG 幾何），不做任何模型計算。
- 敏感度情境定義在 tools/html_scenarios.yaml（值取自 Inputs 的低／高欄、TK_Link 的 Lo／Hi 具名格）。
- 每個數字以 <span class="n" data-ref data-y data-sc data-fmt data-v> 標示來源，title 顯示 Excel 位置；
  tests/test_html.py 以 engine（pycel）逐一比對。
- 取位寫法：具名範圍（多年範圍需指定年度）；「頁!列 ID」（年度欄 D–I＝2025–2030，未指定年度＝D 欄）；
  「頁!列 ID:欄」（指定欄，例 Funding!F88:G＝回流對照表的「模型付款」欄、SRC_ANT!SRC_ANT_373:F＝低欄）。
- 輸出單一檔案、離線可開：CSS 內嵌、圖表為內嵌 SVG、系統字型，無任何外部 URL。

用法：python3 tools/build_html.py [--outdir dist]
"""
from __future__ import annotations

import argparse
import html
import math
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "tests" / "parity"))
from engine import current_model_path, parse_ref  # noqa: E402

YEARS = [2025, 2026, 2027, 2028, 2029, 2030]
YEAR_COL = {y: chr(ord("D") + i) for i, y in enumerate(YEARS)}   # D–I＝2025–2030（各頁共同版面）
SCEN_FILE = REPO / "tools" / "html_scenarios.yaml"
MINUS = "−"
ROWID_RE = re.compile(r"[A-Z]{1,3}\d{1,3}|C\d{2,3}|INP_\d{3}|SRC_ANT_\d{3}")


# ── 取位：具名範圍或「頁!列 ID[:欄]」→（頁, 格, 列 ID）──────────────────────
class Locator:
    """以公式模式的活頁簿建立位址對照（不讀計算值）。"""

    def __init__(self, path: Path):
        wb = openpyxl.load_workbook(path)
        self.wb = wb
        self.names = {k: v.attr_text for k, v in wb.defined_names.items()}
        self.rowid: dict[tuple[str, str], int] = {}       # (頁, 列 ID) → 列號
        self.rowid_of: dict[tuple[str, int], str] = {}    # (頁, 列號) → 列 ID
        for ws in wb:
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if isinstance(v, str) and ROWID_RE.fullmatch(v):
                    self.rowid.setdefault((ws.title, v), r)
                    self.rowid_of[(ws.title, r)] = v
        self.inputs = wb["Inputs"]

    def locate(self, ref: str, year: int | None = None) -> tuple[str, str, str]:
        if "!" in ref:
            sheet, rid = ref.split("!", 1)
            col = None
            if ":" in rid:
                rid, col = rid.split(":", 1)
            if re.fullmatch(r"[A-Z]+\d+", rid) and (sheet, rid) not in self.rowid:   # 直接格位（TK_Link!B4 等）
                return sheet, rid, rid
            r = self.rowid[(sheet, rid)]
            col = col or (YEAR_COL[year] if year is not None else "D")
            return sheet, f"{col}{r}", rid
        sheet, coord = parse_ref(self.names[ref])
        coord = coord.replace("$", "")
        if ":" in coord:
            a, _ = coord.split(":")
            r = int(re.sub(r"[A-Z]+", "", a))
            if year is None:
                raise ValueError(f"{ref} 是多年範圍，需指定年度")
            coord = f"{YEAR_COL[year]}{r}"
        r = int(re.sub(r"[A-Z]+", "", coord))
        return sheet, coord, self.rowid_of.get((sheet, r), "")

    def const(self, name: str):
        """單格具名範圍的常數值（公式模式活頁簿；只用於輸入格）。"""
        sheet, coord = parse_ref(self.names[name])
        return self.wb[sheet][coord.replace("$", "")].value

    # Inputs 的值／低／高欄（常數格）；TK NonNV 快照格的 Lo／Hi 具名格
    def input_cell(self, name: str, which: str):
        if name.startswith("TK_"):
            return self.const(name + {"val": "", "lo": "Lo", "hi": "Hi"}[which])
        sheet, coord = parse_ref(self.names[name])
        coord = coord.replace("$", "")
        r = int(re.sub(r"[A-Z]+", "", coord))
        col = {"val": "E", "lo": "F", "hi": "G"}[which]
        return self.inputs[f"{col}{r}"].value


def load_scenarios(path: Path = SCEN_FILE) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def resolve_overrides(scn_id: str, loc: Locator, cfg: dict | None = None) -> dict[str, float]:
    """情境 → {INP_nnn 或 TK_NNV_*: 值}；值一律取自 Excel（Inputs 低／高欄、TK_Link Lo／Hi）。"""
    cfg = cfg or load_scenarios()
    out: dict[str, float] = {}
    for g in cfg["scenarios"][scn_id]:
        for name, spec in cfg["groups"][g]["set"].items():
            if spec in ("lo", "hi"):
                v = loc.input_cell(name, spec)
            elif isinstance(spec, dict) and "mul" in spec:
                v = loc.input_cell(name, "val") * loc.input_cell(spec["mul"], "val")
            elif isinstance(spec, (int, float)):
                v = spec
            else:
                raise ValueError(f"{g}.{name}: 不支援的寫法 {spec!r}")
            if v is None:
                raise ValueError(f"{g}.{name}: 欄位為空")
            out[name] = v
    return out


# ── 格式化（測試共用）───────────────────────────────────────────────────
def fmt(v, kind: str) -> str:
    if kind == "yr":
        return str(int(round(v)))
    if kind == "yt":                       # 年度或文字（例：首次缺口年「無」）
        return str(int(round(v))) if isinstance(v, (int, float)) else str(v)
    if kind == "pct0":
        x = round(v * 100)
        return f"{MINUS if x < 0 else ''}{abs(x):d}%"
    if kind == "pct1":
        x = round(v * 100, 1)
        return f"{MINUS if x < 0 else ''}{abs(x):.1f}%"
    m = re.fullmatch(r"(s?)f(\d)", kind)  # sfN：正值加「＋」
    if m:
        n = int(m.group(2))
        x = round(v, n)
        if x == 0:
            x = 0.0
        sign = MINUS if x < 0 else ("＋" if (m.group(1) and x > 0) else "")
        return f"{sign}{abs(x):,.{n}f}"
    if kind == "text":
        return str(v)
    raise ValueError(kind)


# ── 值來源（LibreOffice 重算後的 xlsx）──────────────────────────────────
class Values:
    def __init__(self, loc: Locator, books: dict[str, Path]):
        self.loc = loc
        self.wb = {k: openpyxl.load_workbook(p, data_only=True) for k, p in books.items()}
        self.used: list[tuple] = []

    def get(self, ref: str, year: int | None = None, sc: str = "base"):
        sheet, coord, _ = self.loc.locate(ref, year)
        v = self.wb[sc][sheet][coord].value
        if v is None or v == "":
            raise ValueError(f"{sc}:{ref}@{year} 為空（{sheet}!{coord}）")
        if isinstance(v, str) and v.startswith("#"):
            raise ValueError(f"{sc}:{ref}@{year} 錯誤值 {v}")
        return v

    def title(self, ref: str, year: int | None, sc: str) -> str:
        sheet, coord, rid = self.loc.locate(ref, year)
        parts = [f"{sheet} {rid}".strip() if rid else sheet, coord]
        if "!" not in ref:
            parts.append(ref)
        if year is not None:
            parts.append(str(year))
        if sc != "base":
            parts.append(f"情境 {sc}")
        return "｜".join(parts)

    def n(self, ref: str, year: int | None = None, f: str = "f1", sc: str = "base", cls: str = "") -> str:
        """可追溯的數字 span。"""
        v = self.get(ref, year, sc)
        self.used.append((ref, year, sc, f))
        y = "" if year is None else f' data-y="{year}"'
        c = f"n {cls}".strip()
        return (f'<span class="{c}" data-ref="{html.escape(ref)}"{y} data-sc="{sc}" data-fmt="{f}" '
                f'data-v="{html.escape(repr(v))}" title="{html.escape(self.title(ref, year, sc))}">{html.escape(fmt(v, f))}</span>')

    def t(self, ref: str, year: int | None = None, f: str = "f1", sc: str = "base", **attrs) -> str:
        """SVG <text>，同樣帶 data-ref（測試一併比對）。"""
        v = self.get(ref, year, sc)
        self.used.append((ref, year, sc, f))
        y = "" if year is None else f' data-y="{year}"'
        a = " ".join(f'{k.replace("_", "-")}="{val}"' for k, val in attrs.items())
        return (f'<text {a} data-ref="{html.escape(ref)}"{y} data-sc="{sc}" data-fmt="{f}" data-v="{html.escape(repr(v))}">'
                f'<title>{html.escape(self.title(ref, year, sc))}</title>{html.escape(fmt(v, f))}</text>')

    def src(self, *refs: str) -> str:
        """圖表下方小字：Excel 位置（頁＋列 ID，具名範圍附註）。"""
        out = []
        for r in refs:
            multi = "!" not in r and ":" in self.loc.names.get(r, "")
            sheet, _, rid = self.loc.locate(r, YEARS[0] if ("!" in r or multi) else None)
            out.append(f"{sheet} {rid}" + ("" if "!" in r else f"（{r}）"))
        return '<p class="src">Excel：' + "、".join(dict.fromkeys(out)) + "</p>"


# ── SVG 小工具 ─────────────────────────────────────────────────────────
def nice_max(x: float) -> float:
    """座標軸上限：取 {1, 1.2, 1.6, 2, 2.4, 4, 6, 8, 10} × 10^k 中第一個 ≥ x 的值（4 等分刻度皆為整齊值）。"""
    p = 10 ** math.floor(math.log10(x))
    for m in (1, 1.2, 1.6, 2, 2.4, 4, 6, 8, 10):
        if m * p >= x - 1e-12:
            return m * p
    return 10 * p


def nice_range(lo: float, hi: float, n: int = 5) -> tuple[float, float, int]:
    """含正負的座標軸：步長取 {1, 2, 2.5, 5} × 10^k，回傳（下限, 上限, 刻度數）。"""
    raw = (hi - lo) / n
    p = 10 ** math.floor(math.log10(raw))
    step = next(m * p for m in (1, 2, 2.5, 5, 10) if m * p >= raw - 1e-12)
    a, b = math.floor(lo / step) * step, math.ceil(hi / step) * step
    return a, b, int(round((b - a) / step))


def axis(W, H, L, R, T, B, ymin, ymax, ticks, unit=""):
    out = []
    for i in range(ticks + 1):
        val = ymin + (ymax - ymin) * i / ticks
        y = T + (H - T - B) * (1 - (val - ymin) / (ymax - ymin))
        out.append(f'<line class="grid{" zero" if abs(val) < 1e-9 else ""}" x1="{L}" x2="{W - R}" y1="{y:.1f}" y2="{y:.1f}"/>')
        lab = f"{MINUS if val < -1e-9 else ''}{abs(val):g}"
        out.append(f'<text class="tick" x="{L - 6}" y="{y + 4:.1f}" text-anchor="end">{lab}</text>')
    if unit:
        out.append(f'<text class="tick" x="2" y="{T - 10}">{unit}</text>')
    return out


def ymap(val, H, T, B, ymin, ymax):
    return T + (H - T - B) * (1 - (val - ymin) / (ymax - ymin))


# ── 頁面 ───────────────────────────────────────────────────────────────
CSS = r"""
:root{--bg:#fbfaf7;--fg:#1d2329;--muted:#5d6670;--line:#d9dde1;--card:#ffffff;--accent:#0d5c63;
--rev:#2563a8;--cost:#c2410c;--gap:#b42318;--pos:#2f855a;--sub:#2563a8;--seat:#7fb2e5;--api:#5aa0d8;--ads:#a3c9ea;--net:#1d2329;--tgt:#7a3e9d;
--inf:#2f855a;--rd:#9ae6b4;--con:#4a5568;--own:#a0aec0;--amz:#d69e2e;--oai:#7a3e9d;--warn-bg:#fff4e5;--ok:#2f855a}
@media (prefers-color-scheme: dark){:root{--bg:#14181c;--fg:#e8ebee;--muted:#9aa4ae;--line:#323a42;--card:#1b2127;--accent:#5fc4cb;
--rev:#6aa7ea;--cost:#f08a4b;--gap:#ff7b72;--pos:#56c288;--sub:#6aa7ea;--seat:#3f6f9f;--api:#3f7fbf;--ads:#9cc7ee;--net:#e8ebee;--tgt:#c792ea;
--inf:#56c288;--rd:#2e6b4a;--con:#a0aec0;--own:#5a6573;--amz:#e6b450;--oai:#c792ea;--warn-bg:#2b2416;--ok:#56c288}}
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%}
body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI","PingFang TC","Microsoft JhengHei","Noto Sans CJK TC","Noto Sans TC","Heiti TC",sans-serif;font-variant-numeric:tabular-nums}
main{max-width:980px;margin:0 auto;padding:24px 16px 48px}
header h1{font-size:22px;line-height:1.3;margin:0 0 4px}
header .meta{color:var(--muted);font-size:13px}
section{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:18px 18px 14px;margin:16px 0}
h2{font-size:17px;margin:0 0 10px;display:flex;gap:8px;align-items:baseline}
h2 .no{color:var(--accent);font-weight:700}
h3{font-size:15px;margin:14px 0 6px}
p{margin:6px 0}
.lead{font-size:16px;font-weight:600;border-left:4px solid var(--gap);padding:6px 0 6px 12px;margin:4px 0 14px}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:10px}
.kpi{border:1px solid var(--line);border-radius:8px;padding:10px 12px}
.kpi .k{font-size:12px;color:var(--muted)}
.kpi .v{font-size:24px;font-weight:700;line-height:1.25}
.kpi .s{font-size:12px;color:var(--muted)}
.neg{color:var(--gap)}
.pos{color:var(--pos)}
.n{white-space:nowrap}
.src{font-size:11.5px;color:var(--muted);margin:4px 0 0}
.note{font-size:13px;color:var(--muted)}
.cw{overflow-x:auto;-webkit-overflow-scrolling:touch}
.tw{overflow-x:auto;-webkit-overflow-scrolling:touch;margin:8px 0}
table{border-collapse:collapse;width:100%;font-size:13px}
th,td{padding:5px 6px;border-bottom:1px solid var(--line);text-align:right;white-space:nowrap}
th:first-child,td:first-child{text-align:left;white-space:normal;min-width:9em}
thead th{color:var(--muted);font-weight:600;border-bottom:1.5px solid var(--muted)}
tr.em td{font-weight:700}
tr.sep td{border-top:1.5px solid var(--muted)}
svg{width:100%;height:auto;display:block}
svg text{fill:var(--fg);font-size:12px;font-family:inherit}
svg .tick{fill:var(--muted);font-size:11px}
svg .grid{stroke:var(--line);stroke-width:1}
svg .grid.zero{stroke:var(--muted)}
svg .lbl{font-size:11.5px}
svg .gaplbl{fill:var(--gap);font-weight:700;font-size:12px}
svg .poslbl{fill:var(--pos);font-weight:700;font-size:12px}
.legend{display:flex;flex-wrap:wrap;gap:4px 14px;font-size:12.5px;color:var(--muted);margin:2px 0 4px}
.legend i{display:inline-block;width:11px;height:11px;border-radius:2px;margin-right:5px;vertical-align:-1px}
.must{background:var(--warn-bg);border-radius:8px;padding:10px 12px;margin:10px 0}
ol,ul{padding-left:1.3em;margin:6px 0}
li{margin:4px 0}
.tag{display:inline-block;font-size:11.5px;border:1px solid var(--line);border-radius:4px;padding:0 5px;margin-right:4px;color:var(--muted)}
dl{display:grid;grid-template-columns:max-content 1fr;gap:4px 14px;font-size:13px;margin:6px 0}
dt{color:var(--muted)}dd{margin:0;word-break:break-all}
footer{color:var(--muted);font-size:12px;text-align:center;margin-top:20px}
@media (max-width:640px){
 main{padding:16px 12px 40px}
 .kpis{grid-template-columns:1fr 1fr}
 .kpi .v{font-size:20px}
 section{padding:14px 12px 10px}
 header h1{font-size:19px}
 table{font-size:11.5px}
 th,td{padding:4px 3px}
 th:first-child,td:first-child{min-width:6.5em}
 .cw svg{min-width:560px}
}
@media print{section{break-inside:avoid}}
"""


def chart_gw(V: Values) -> str:
    """每 VR 等值 GW 營收 vs 全成本。2025 分母極小（數值為其他年的 5–15 倍），軸以 2026–2030 定上限，2025 長條截斷並標實值。"""
    W, H, L, R, T, B = 640, 300, 44, 10, 26, 34
    rev = [V.get("COST_PropRev_VR", y) for y in YEARS]
    full = [V.get("COST_PropFull_VR", y) for y in YEARS]
    ymax = nice_max(max(full[1:] + rev[1:]) * 1.12)
    o = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="每 VR 等值 GW 營收與全成本 2025–2030">']
    o += axis(W, H, L, R, T, B, 0, ymax, 4, "$B/GW/年")
    slot = (W - L - R) / len(YEARS)
    bw = slot * 0.3
    for i, y in enumerate(YEARS):
        x0 = L + slot * i + slot * 0.18
        clipped = False
        for j, (vals, cls, ref) in enumerate(((rev, "rev", "COST_PropRev_VR"), (full, "cost", "COST_PropFull_VR"))):
            v = vals[i]
            yt = ymap(min(v, ymax), H, T, B, 0, ymax)
            o.append(f'<rect x="{x0 + j * bw:.1f}" y="{yt:.1f}" width="{bw - 2:.1f}" height="{H - B - yt:.1f}" style="fill:var(--{cls})"/>')
            if v > ymax:
                clipped = True
                xb = x0 + j * bw
                o.append(f'<path d="M{xb - 1:.1f},{T + 14:.1f} l{bw:.1f},-6 v5 l-{bw:.1f},6 z" style="fill:var(--card)"/>')
                o.append(V.t(ref, y, "f0", x=f"{xb + bw / 2 - 1:.1f}", y=f"{T + 34 + 14 * j:.1f}", text_anchor="middle",
                             **{"class": "lbl"}, style="fill:#fff;font-weight:700"))
        top = ymap(min(max(full[i], rev[i]), ymax), H, T, B, 0, ymax)
        g = V.get("COST_PropGap_VR", y)
        o.append(V.t("COST_PropGap_VR", y, "sf1", x=f"{x0 + bw:.1f}", y=f"{(T - 4 if clipped else top - 6):.1f}", text_anchor="middle",
                     **{"class": "gaplbl" if g < 0 else "poslbl"}))
        o.append(f'<text class="tick" x="{L + slot * i + slot / 2:.1f}" y="{H - 12}" text-anchor="middle">{y}</text>')
    o.append("</svg>")
    return "\n".join(o)


def chart_rev(V: Values) -> str:
    W, H, L, R, T, B = 640, 300, 44, 10, 26, 34
    gross = [V.get("REV_GrossCapped", y) for y in YEARS]
    ymax = nice_max(max(gross) * 1.1)
    o = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="營收結構（個人訂閱、企業席位、API、其他）與管理層目標">']
    o += axis(W, H, L, R, T, B, 0, ymax, 4, "$B")
    slot = (W - L - R) / len(YEARS)
    bw = slot * 0.36
    for i, y in enumerate(YEARS):
        xc = L + slot * i + slot * 0.42
        base = 0.0
        for ref, cls in (("REV_SubConsumer", "sub"), ("REV_SubSeats", "seat"), ("REV_API", "api"), ("REV_Other", "ads")):
            v = V.get(ref, y)
            if v <= 0:
                continue
            y1, y0 = ymap(base + v, H, T, B, 0, ymax), ymap(base, H, T, B, 0, ymax)
            o.append(f'<rect x="{xc - bw / 2:.1f}" y="{y1:.1f}" width="{bw:.1f}" height="{max(y0 - y1, 0):.1f}" style="fill:var(--{cls})"/>')
            base += v
        net = V.get("REV_NetCapped", y)
        yn = ymap(net, H, T, B, 0, ymax)
        o.append(f'<line x1="{xc - bw / 2 - 4:.1f}" x2="{xc + bw / 2 + 4:.1f}" y1="{yn:.1f}" y2="{yn:.1f}" style="stroke:var(--net);stroke-width:2.5"/>')
        o.append(V.t("REV_GrossCapped", y, "f1", x=f"{xc:.1f}", y=f"{ymap(base, H, T, B, 0, ymax) - 5:.1f}", text_anchor="middle", **{"class": "lbl"}))
        if 2026 <= y <= 2029:
            yt = ymap(V.get("RVS_Target", y), H, T, B, 0, ymax)
            xd = xc + bw / 2 + 8
            o.append(f'<path d="M{xd:.1f},{yt - 6:.1f} l6,6 l-6,6 l-6,-6 z" style="fill:var(--tgt)"/>')
            o.append(V.t("RVS_Target", y, "f0", x=f"{xd + 9:.1f}", y=f"{yt + 4:.1f}", **{"class": "lbl"}, style="fill:var(--tgt)"))
        o.append(f'<text class="tick" x="{xc:.1f}" y="{H - 12}" text-anchor="middle">{y}</text>')
    o.append("</svg>")
    return "\n".join(o)


# run-rate 里程碑（D23；只對照）：SRC_ANT 列 ID（每月一筆；同月重複報導取一筆）。月份位置取自 SRC_ANT 的「期間」欄。
RUNRATE_IDS = ["SRC_ANT_008", "SRC_ANT_011", "SRC_ANT_012", "SRC_ANT_013", "SRC_ANT_015", "SRC_ANT_016", "SRC_ANT_017",
               "SRC_ANT_019", "SRC_ANT_020", "SRC_ANT_021", "SRC_ANT_022", "SRC_ANT_023"]
RUNRATE_RANGE_ID = "SRC_ANT_373"      # 2026 年底投資人預期（只有低、高）


def period_pos(loc: Locator, sid: str) -> tuple[float, str]:
    """SRC_ANT 期間欄（I）→ 時間軸位置（年＋月份比例）；月初＝月首、月底＝月末、其餘＝月中。"""
    r = loc.rowid[("SRC_ANT", sid)]
    txt = str(loc.wb["SRC_ANT"][f"I{r}"].value)
    m = re.match(r"(\d{4})-(\d{2})", txt)
    yr, mo = int(m.group(1)), int(m.group(2))
    off = 0.0 if "月初" in txt else (1.0 if "月底" in txt else 0.5)
    return yr + (mo - 1 + off) / 12, f"{yr}-{mo:02d}"


def chart_runrate(V: Values) -> str:
    W, H, L, R, T, B = 640, 280, 44, 10, 26, 34
    x0, x1 = 2024.75, 2027.0
    hi = V.get(f"SRC_ANT!{RUNRATE_RANGE_ID}:G")
    ymax = nice_max(hi * 1.08)
    X = lambda t: L + (W - L - R) * (t - x0) / (x1 - x0)  # noqa: E731
    o = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="年化營收里程碑與模型年營收">']
    o += axis(W, H, L, R, T, B, 0, ymax, 4, "$B（年化／年）")
    for yr in (2025, 2026, 2027):
        o.append(f'<line class="grid" x1="{X(yr):.1f}" x2="{X(yr):.1f}" y1="{T}" y2="{H - B}"/>')
    for yr in (2025, 2026):
        o.append(f'<text class="tick" x="{X(yr + 0.5):.1f}" y="{H - 12}" text-anchor="middle">{yr}</text>')
        v = V.get("REV_GrossCapped", yr)
        yy = ymap(v, H, T, B, 0, ymax)
        o.append(f'<rect x="{X(yr) + 2:.1f}" y="{yy:.1f}" width="{X(yr + 1) - X(yr) - 4:.1f}" height="{H - B - yy:.1f}" style="fill:var(--rev);opacity:.22"/>')
        o.append(f'<line x1="{X(yr) + 2:.1f}" x2="{X(yr + 1) - 2:.1f}" y1="{yy:.1f}" y2="{yy:.1f}" style="stroke:var(--rev);stroke-width:2.5"/>')
        o.append(V.t("REV_GrossCapped", yr, "f1", x=f"{X(yr + 0.25):.1f}", y=f"{yy - 6:.1f}", text_anchor="middle", **{"class": "lbl"}, style="fill:var(--rev);font-weight:700"))
    pts = []
    for sid in RUNRATE_IDS:
        t, _ = period_pos(V.loc, sid)
        pts.append((X(t), ymap(V.get(f"SRC_ANT!{sid}:E"), H, T, B, 0, ymax), sid))
    o.append('<polyline fill="none" style="stroke:var(--cost);stroke-width:1.5" points="' + " ".join(f"{x:.1f},{y:.1f}" for x, y, _ in pts) + '"/>')
    for k, (x, y, sid) in enumerate(pts):
        o.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.5" style="fill:var(--cost)"/>')
        if k >= 9:   # 2026-04 起標值（其餘數值小，見明細表）
            o.append(V.t(f"SRC_ANT!{sid}:E", None, "f0", x=f"{x + 7:.1f}", y=f"{y + 4:.1f}", text_anchor="start", **{"class": "lbl"}, style="fill:var(--cost)"))
    t, _ = period_pos(V.loc, RUNRATE_RANGE_ID)
    xr = X(min(t, x1 - 0.03))
    ylo, yhi = ymap(V.get(f"SRC_ANT!{RUNRATE_RANGE_ID}:F"), H, T, B, 0, ymax), ymap(hi, H, T, B, 0, ymax)
    o.append(f'<line x1="{xr:.1f}" x2="{xr:.1f}" y1="{ylo:.1f}" y2="{yhi:.1f}" style="stroke:var(--tgt);stroke-width:5;stroke-linecap:round"/>')
    o.append(V.t(f"SRC_ANT!{RUNRATE_RANGE_ID}:F", None, "f0", x=f"{xr - 7:.1f}", y=f"{ylo + 4:.1f}", text_anchor="end", **{"class": "lbl"}, style="fill:var(--tgt)"))
    o.append(V.t(f"SRC_ANT!{RUNRATE_RANGE_ID}:G", None, "f0", x=f"{xr - 7:.1f}", y=f"{yhi + 4:.1f}", text_anchor="end", **{"class": "lbl"}, style="fill:var(--tgt)"))
    o.append("</svg>")
    return "\n".join(o)


def chart_compute(V: Values) -> str:
    W, H, L, R, T, B = 640, 300, 44, 10, 26, 34
    sup = [V.get("CMP_SupplyGW", y) for y in YEARS]
    ymax = nice_max(max(sup) * 1.1)
    o = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="算力需求與供給 GW">']
    o += axis(W, H, L, R, T, B, 0, ymax, 4, "GW")
    slot = (W - L - R) / len(YEARS)
    bw = slot * 0.32
    for i, y in enumerate(YEARS):
        x0 = L + slot * i + slot * 0.16
        stacks = ((("CMP_InfGW_Eff", "inf"), ("CMP_RDGW", "rd"), "CMP_DemandGW"),
                  (("CMP_SupplyContractGW", "con"), ("CMP_Supply_Owned", "own"), "CMP_SupplyGW"))
        for j, (a, b, tot) in enumerate(stacks):
            base = 0.0
            x = x0 + j * (bw + 2)
            for ref, cls in (a, b):
                v = V.get(ref, y)
                y1, y0 = ymap(base + v, H, T, B, 0, ymax), ymap(base, H, T, B, 0, ymax)
                if y0 - y1 > 0.05:
                    o.append(f'<rect x="{x:.1f}" y="{y1:.1f}" width="{bw:.1f}" height="{y0 - y1:.1f}" style="fill:var(--{cls})"/>')
                base += v
            o.append(V.t(tot, y, "f1", x=f"{x + bw / 2:.1f}", y=f"{ymap(base, H, T, B, 0, ymax) - 5:.1f}", text_anchor="middle", **{"class": "lbl"}))
        o.append(f'<text class="tick" x="{L + slot * i + slot / 2:.1f}" y="{H - 12}" text-anchor="middle">{y}</text>')
    o.append("</svg>")
    return "\n".join(o)


def chart_cash(V: Values) -> str:
    """年底現金：基準 vs S7 合併不利情境，最低現金以短橫線標示。"""
    W, H, L, R, T, B = 640, 280, 44, 10, 26, 34
    yrs = YEARS[1:]
    series = (("FND_CashEnd", "rev"), ("Funding!F79", "cost"))
    ymax = nice_max(max(V.get(r, y) for r, _ in series for y in yrs) * 1.12)
    o = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="年底現金：基準與 S7 合併不利情境">']
    o += axis(W, H, L, R, T, B, 0, ymax, 4, "$B")
    slot = (W - L - R) / len(yrs)
    bw = slot * 0.3
    for i, y in enumerate(yrs):
        x0 = L + slot * i + slot * 0.18
        for j, (ref, cls) in enumerate(series):
            v = V.get(ref, y)
            yt = ymap(v, H, T, B, 0, ymax)
            o.append(f'<rect x="{x0 + j * bw:.1f}" y="{yt:.1f}" width="{bw - 2:.1f}" height="{H - B - yt:.1f}" style="fill:var(--{cls})"/>')
            o.append(V.t(ref, y, "f0", x=f"{x0 + j * bw + bw / 2 - 1:.1f}", y=f"{yt - 5:.1f}", text_anchor="middle", **{"class": "lbl"}))
        ym = ymap(V.get("FND_MinCash", y), H, T, B, 0, ymax)
        o.append(f'<line x1="{x0 - 4:.1f}" x2="{x0 + 2 * bw + 2:.1f}" y1="{ym:.1f}" y2="{ym:.1f}" style="stroke:var(--gap);stroke-width:2;stroke-dasharray:4 3"/>')
        o.append(f'<text class="tick" x="{L + slot * i + slot / 2:.1f}" y="{H - 12}" text-anchor="middle">{y}</text>')
    o.append("</svg>")
    return "\n".join(o)


def chart_pair(V: Values) -> str:
    """與 OpenAI v0.6 並排：每 VR 等值 GW 差額 2026–2030（2025 分母效應，兩家皆略）。"""
    W, H, L, R, T, B = 640, 260, 44, 10, 26, 34
    yrs = YEARS[1:]
    series = (("COST_PropGap_VR", "rev", "Anthropic"), ("OAI_COST_PropGap_VR", "oai", "OpenAI"))
    vals = [V.get(r, y) for r, _, _ in series for y in yrs]
    ymin, ymax, ticks = nice_range(min(vals + [0]) * 1.05, max(vals + [0]) * 1.15)
    o = [f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="每 VR 等值 GW 差額：Anthropic 與 OpenAI">']
    o += axis(W, H, L, R, T, B, ymin, ymax, ticks, "$B/GW/年")
    slot = (W - L - R) / len(yrs)
    for ref, cls, _ in series:
        pts = [(L + slot * i + slot / 2, ymap(V.get(ref, y), H, T, B, ymin, ymax)) for i, y in enumerate(yrs)]
        o.append(f'<polyline fill="none" style="stroke:var(--{cls});stroke-width:2.5" points="' + " ".join(f"{x:.1f},{y:.1f}" for x, y in pts) + '"/>')
        for (x, yy), y in zip(pts, yrs):
            o.append(f'<circle cx="{x:.1f}" cy="{yy:.1f}" r="4" style="fill:var(--{cls})"/>')
            above = cls == "rev"
            o.append(V.t(ref, y, "sf1", x=f"{x:.1f}", y=f"{yy - 9 if above else yy + 17:.1f}", text_anchor="middle", **{"class": "lbl"}, style=f"fill:var(--{cls});font-weight:700"))
    for i, y in enumerate(yrs):
        o.append(f'<text class="tick" x="{L + slot * i + slot / 2:.1f}" y="{H - 12}" text-anchor="middle">{y}</text>')
    o.append("</svg>")
    return "\n".join(o)


def year_table(V: Values, rows: list[tuple], years=YEARS, head="$B") -> str:
    """rows：(標籤, ref, fmt, 樣式)。"""
    h = "".join(f"<th>{y}</th>" for y in years)
    o = [f'<div class="tw"><table><thead><tr><th>{head}</th>{h}</tr></thead><tbody>']
    for lab, ref, f, cls in rows:
        cells = []
        for y in years:
            try:
                cells.append(f"<td>{V.n(ref, y, f)}</td>")
            except ValueError:
                cells.append("<td>—</td>")
        o.append(f'<tr class="{cls}"><td>{lab}</td>{"".join(cells)}</tr>')
    o.append("</tbody></table></div>")
    return "\n".join(o)


def cell(V: Values, ref: str, year=None, f="f1", sc="base") -> str:
    try:
        return V.n(ref, year, f, sc)
    except ValueError:
        return "—"


def build(outdir: Path) -> tuple[Path, Path, Values]:
    from parity_lib import lo_recalc, set_inputs

    model = current_model_path()
    loc = Locator(model)
    cfg = load_scenarios()
    tmp = Path(tempfile.mkdtemp(prefix="a5_html_"))
    books = {}
    for sc in cfg["scenarios"]:
        src = tmp / f"in_{sc}" / model.name
        src.parent.mkdir(parents=True)
        ov = resolve_overrides(sc, loc, cfg)
        if ov:
            set_inputs(model, src, ov)
        else:
            shutil.copy(model, src)
        books[sc] = lo_recalc(src, tmp / f"out_{sc}")
    V = Values(loc, books)
    if V.get("CHK_Errors", None, "base") != 0:
        raise SystemExit("基準 CHK_Errors ≠ 0，不產生成品")

    stamp = model.name.split("_")[0]
    ver = re.search(r"_v([\d.]+)\.xlsx$", model.name).group(1)
    html_path = outdir / f"{stamp}_Anthropic收支模型_v{ver.replace('.', '_')}.html"
    xlsx_path = outdir / f"{stamp}_Anthropic收支模型_v{ver.replace('.', '_')}.xlsx"
    page = render(V, model, ver, stamp)
    outdir.mkdir(parents=True, exist_ok=True)
    html_path.write_text(page, encoding="utf-8")
    shutil.copyfile(model, xlsx_path)
    shutil.rmtree(tmp, ignore_errors=True)
    return html_path, xlsx_path, V


def src_meta(model: Path) -> tuple[str, int, int]:
    """SRC_ANT 最新文件日期、列數、來源引用 IPO 招股書（草稿）者的列數（只用於資料說明，不進任何數字）。"""
    ws = openpyxl.load_workbook(model, read_only=True)["SRC_ANT"]
    ds, n, pros = [], 0, 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not (isinstance(row[0], str) and row[0].startswith("SRC_ANT_")):
            continue
        n += 1
        d = row[13]
        if isinstance(d, str) and re.fullmatch(r"\d{4}-\d{2}(-\d{2})?", d.strip()):
            ds.append(d.strip())
        s = f"{row[11] or ''} {row[22] or ''}"
        if "招股書" in s or "prospectus" in s.lower() or "說明書" in s:
            pros += 1
    return (max(ds) if ds else "—"), n, pros


def render(V: Values, model: Path, ver: str, stamp: str) -> str:
    n = V.n
    Y = 2030
    date_disp = f"{stamp[:4]}-{stamp[4:6]}-{stamp[6:]}"
    tk_file, tk_ver, tk_sha, tk_date = (V.get(f"TK_Link!B{i}") for i in (3, 4, 5, 6))
    oai_file, oai_sha, oai_commit = V.get("OAI_File"), V.get("OAI_SHA256"), V.get("OAI_Commit")
    src_date, src_n, src_pros = src_meta(model)
    chk_ids = sorted((k[1] for k in V.loc.rowid if k[0] == "Checks"), key=lambda s: int(s[1:]))
    report = f"docs/reports/{stamp}_v{ver}.md"

    S = []
    # ── 標頭 ──
    S.append(f"""<header><h1>Anthropic 收支模型 v{ver}｜一頁摘要</h1>
<div class="meta">{date_disp}｜Excel <code>{html.escape(model.name)}</code>｜Tokenomics {html.escape(str(tk_ver))}（<code>{html.escape(str(tk_sha))[:7]}</code>）｜單位 $B（十億美元），曆年制｜數字游標停留可見 Excel 位置；手機上圖表可左右滑動｜<b>2025–2026 財務多數來自未公開的招股書草稿（Reuters 等轉述），屬利害關係方說法</b></div></header>""")

    # ── ① 命題與結論 ──
    S.append(f"""<section id="s1"><h2><span class="no">①</span>命題與一句話結論</h2>
<p class="note">命題（與 OpenAI v0.6 同一句）：Anthropic 每 VR 等值 GW 的年營收，能否覆蓋每 GW 年全成本？若不能，缺口要多少外部資金、由誰補？</p>
<p class="lead">2026–2028 年還不能（每 VR 等值 GW 差額 {n("COST_PropGap_VR", 2026)}／{n("COST_PropGap_VR", 2027)}／{n("COST_PropGap_VR", 2028)} $B/GW/年），2029 年起轉正（2030 年 {n("COST_PropGap_VR", Y, "sf1")}，覆蓋率 {n("COST_Coverage", Y, "pct0")}）；
2026–2028 的現金缺口已由 2026 年交割的 ${n("FND_Committed", 2026, "f0")}B 股權補足，2026–2030 累計外部資金需求 ${n("FND_ExtNeedCum", Y, "f0")}B。
<b>要成立必須</b>：2029 年後不再大量新增算力合約（模型只計已簽合約），且 API 需求成長與價格反應不同時落在區間下緣。</p>
<div class="kpis">
<div class="kpi"><div class="k">2030 每 VR 等值 GW：營收 vs 全成本</div><div class="v">{n("COST_PropRev_VR", Y)} <span class="note">vs</span> {n("COST_PropFull_VR", Y)}</div><div class="s">$B/GW/年；差額 <b class="pos">{n("COST_PropGap_VR", Y, "sf1")}</b></div></div>
<div class="kpi"><div class="k">覆蓋率（營收淨額 ÷ 全成本）</div><div class="v">{n("COST_Coverage", Y, "pct0")}</div><div class="s">2030；2026–2028 為 {n("COST_Coverage", 2026, "pct0")}／{n("COST_Coverage", 2027, "pct0")}／{n("COST_Coverage", 2028, "pct0")}；全成本含股權報酬、現金口徑</div></div>
<div class="kpi"><div class="k">累計外部資金需求 2026–2030</div><div class="v">{n("FND_ExtNeedCum", Y, "f0")}</div><div class="s">$B；OpenAI v0.6 同期 <b class="neg">{n("OAI_FND_ExtNeedCum", Y, "f1")}</b></div></div>
<div class="kpi"><div class="k">只靠已到位融資的現金谷底</div><div class="v">{n("FND_CashNoExtTrough", None, "f1")}</div><div class="s">$B，2028 年；當年最低現金 {n("FND_MinCash", 2028, "f1")}</div></div>
</div>
<p class="note">翻轉結論需要兩個需求驅動同時落在下緣：API 任務數成長取低＋價格彈性取弱端，2030 差額 {n("COST_PropGap_VR", Y, "sf1", "combo1")}、累計外部資金需求 {n("FND_ExtNeedCum", Y, "f1", "combo1")}（見 ⑥）。用管理層較舊的營收目標、同一支出，反向累計外部資金需求 {n("RVS_ExtNeedCum", Y, "f1")}（只作對照）。</p>
</section>""")

    # ── ② 每 VR 等值 GW ──
    S.append(f"""<section id="s2"><h2><span class="no">②</span>每 VR 等值 GW：營收 vs 全成本（2025–2030）</h2>
<div class="legend"><span><i style="background:var(--rev)"></i>營收淨額（截頂後；扣雲端平台抽成）</span><span><i style="background:var(--cost)"></i>全成本（算力＋非算力＋股權報酬，現金口徑）</span><span><b style="color:var(--gap)">紅</b>／<b style="color:var(--pos)">綠</b>字＝差額</span><span>2025 長條截斷，白字為實值</span></div>
<div class="cw">{chart_gw(V)}</div>
{year_table(V, [("營收淨額", "COST_PropRev_VR", "f1", ""), ("　算力成本（合約實付＋自建資本支出）", "COST_PropCompute_VR", "f1", ""),
                ("　非算力成本（不含股權報酬）", "COST_PropNonComp_VR", "f1", ""), ("　股權報酬", "COST_PropSBC_VR", "f1", ""),
                ("全成本", "COST_PropFull_VR", "f1", ""), ("差額（命題 1；現金・每 VR 等值 GW）", "COST_PropGap_VR", "sf1", "em"),
                ("覆蓋率", "COST_Coverage", "pct0", ""), ("分母：供給 VR 等值 GW", "CMP_Supply_VReq", "f2", ""),
                ("對照：經濟口徑差額（自建改以持有成本年化）", "COST_PropGapEcon_VR", "sf1", "sep"), ("對照：每實體 GW 差額", "COST_PropGap_GW", "sf1", ""),
                ("分母：供給實體 GW", "CMP_SupplyGW", "f2", ""), ("差額總額（$B）", "COST_GapCash", "sf1", "")], head="$B/GW/年")}
<p class="note">2025 分母只有 {n("CMP_Supply_VReq", 2025, "f2")} VR 等值 GW（2025 機隊為 Trainium2、TPU v6e、Hopper，換成 VR200 等值只有實體 {n("CMP_SupplyGW", 2025, "f2")} GW 的 6%），每 GW 數字因此偏大，請看每實體 GW 列（{n("COST_PropGap_GW", 2025, "sf1")}）。三種口徑轉正年同為 2029；2026 經濟口徑為正（{n("COST_PropGapEcon_VR", 2026, "sf1")}）是因自建資本支出在現金口徑一次認列。</p>
{V.src("COST_PropRev_VR", "COST_PropCompute_VR", "COST_PropNonComp_VR", "COST_PropSBC_VR", "COST_PropFull_VR", "COST_PropGap_VR", "COST_Coverage", "CMP_Supply_VReq", "COST_PropGapEcon_VR", "COST_PropGap_GW", "CMP_SupplyGW", "COST_GapCash")}
</section>""")

    # ── ③ 營收結構與 run-rate ──
    rr_rows = []
    for sid in RUNRATE_IDS:
        _, lab = period_pos(V.loc, sid)
        rr_rows.append(f"<tr><td>{lab}</td><td>{n(f'SRC_ANT!{sid}:E', None, 'f0')}</td><td>{html.escape(sid)}</td></tr>")
    rr_rows.append(f"<tr><td>{period_pos(V.loc, RUNRATE_RANGE_ID)[1]}（投資人預期）</td><td>{n(f'SRC_ANT!{RUNRATE_RANGE_ID}:F', None, 'f0')}–{n(f'SRC_ANT!{RUNRATE_RANGE_ID}:G', None, 'f0')}</td><td>{RUNRATE_RANGE_ID}</td></tr>")
    S.append(f"""<section id="s3"><h2><span class="no">③</span>營收結構、雲端平台抽成、年化營收里程碑</h2>
<div class="legend"><span><i style="background:var(--sub)"></i>個人訂閱（Pro、Max）</span><span><i style="background:var(--seat)"></i>企業席位（Team、Enterprise）</span><span><i style="background:var(--api)"></i>API（含經雲端市集）</span><span><i style="background:var(--ads)"></i>其他</span><span><i style="background:var(--net);height:3px"></i>淨額（扣雲端平台抽成）</span><span><i style="background:var(--tgt);transform:rotate(45deg)"></i>管理層目標（舊版）</span></div>
<div class="cw">{chart_rev(V)}</div>
{year_table(V, [("個人訂閱（Pro、Max）", "REV_SubConsumer", "f1", ""), ("企業席位訂閱（Team、Enterprise）", "REV_SubSeats", "f1", ""), ("API", "REV_API", "f1", ""), ("其他", "REV_Other", "f2", ""),
                ("總額（報導口徑；截頂後）", "REV_GrossCapped", "f1", "em"), ("雲端平台抽成（截頂後）", "REV_PartnerShareCapped", "f1", ""), ("淨額（命題與現金流用）", "REV_NetCapped", "f1", "em"),
                ("企業占（個人＋企業）", "Revenue!R60", "pct0", "sep"), ("　其中 API 經雲端通路的營收", "REV_Channel", "f1", ""),
                ("管理層目標（只對照）", "RVS_Target", "f1", ""), ("差距（目標 − 正向總額）", "RVS_Gap", "f1", "")])}
<p class="note">報導營收為總額（AWS Bedrock、Google Vertex 等雲端市集以總額認列），本模型扣掉平台抽成（2026 起 {n("Revenue!R55", 2026, "pct0")}）後的淨額才進命題與現金流；2030 年抽成 {n("REV_PartnerShareCapped", Y, "f1")}。廣告＝0（D3）。容量上限係數各年皆為 {n("CMP_CapFactor", Y, "f0")}（未受限），所以各線未截頂值＝截頂值。管理層目標都早於 2026 年的實際爆發，全部低於模型正向值。</p>
<h3>模型年營收 vs 年化營收（run-rate）里程碑（D23：只對照，不校準）</h3>
<div class="legend"><span><i style="background:var(--rev)"></i>模型年營收總額（年平均水準）</span><span><i style="background:var(--cost);border-radius:50%"></i>報導年化營收</span><span><i style="background:var(--tgt)"></i>2026 年底投資人預期區間</span></div>
<div class="cw">{chart_runrate(V)}</div>
<div class="tw"><table><thead><tr><th>對照</th><th>2025</th><th>2026</th></tr></thead><tbody>
<tr class="em"><td>模型年營收總額</td><td>{n("REV_GrossCapped", 2025)}</td><td>{n("REV_GrossCapped", 2026)}</td></tr>
<tr><td>由 run-rate 推估的年營收（2025：年初、年底年化的對數平均；2026：上半年實際＋下半年 65 → 年底預期中點）</td><td>{n("Revenue!R72", 2025)}</td><td>{n("REV_RREst2026", None)}</td></tr>
<tr><td>模型 ÷ 推估 − 1</td><td>{n("Revenue!R73", 2025, "pct0")}</td><td>{n("REV_RRGap2026", None, "pct0")}</td></tr>
</tbody></table></div>
<details><summary class="note">年化營收里程碑明細（SRC_ANT；公司新聞稿與媒體轉述，多為利害關係方說法，含第三方雲端通路總額）</summary>
<div class="tw"><table><thead><tr><th>月份</th><th>年化營收 $B</th><th>來源列</th></tr></thead><tbody>{"".join(rr_rows)}</tbody></table></div></details>
<p class="note">2025 模型以認列營收 {n("SRC_ANT_001", None, "f2")} 校準，高於 run-rate 推估 {n("Revenue!R73", 2025, "pct0")}；2026 模型為「上半年實際＋下半年持平於 2026-07 年化 65」，低於里程碑推估 {n("REV_RRGap2026", None, "pct0")}（若年底真到 100–120，2026 會被低估）。OpenAI 內部備忘錄估 Anthropic 2026-04 年化以淨額計只有 {n("SRC_ANT!SRC_ANT_035:E", None, "f0")}（對公司說的 {n("SRC_ANT!SRC_ANT_021:E", None, "f0")}），差在雲端總額認列。</p>
{V.src("REV_SubConsumer", "REV_SubSeats", "REV_API", "REV_Other", "REV_GrossCapped", "REV_PartnerShareCapped", "REV_NetCapped", "Revenue!R60", "REV_Channel", "RVS_Target", "RVS_Gap", "Revenue!R55", "Revenue!R72", "Revenue!R73", "REV_RREst2026", "REV_RRGap2026", "SRC_ANT_001")}
</section>""")

    # ── ④ 算力 ──
    contracts = [  # (標籤, 供給 GW 具名, 實付具名, 模型合約總額列, 招股書列, 差距列)
        ("Google Cloud TPU", "CMP_Supply_Google", "COST_Pay_Google", "Cost!K121", "Cost!K133", "Cost!K143"),
        ("Google／Broadcom 次世代 TPU", "CMP_Supply_Broadcom", "COST_Pay_Broadcom", "Cost!K122", "Cost!K136", "Cost!K146"),
        ("Amazon（Trainium）", "CMP_Supply_Aws", "COST_Pay_Aws", "Cost!K123", "Cost!K134", "Cost!K144"),
        ("Microsoft Azure＋NVIDIA", "CMP_Supply_Azure", "COST_Pay_Azure", "Cost!K124", "Cost!K135", "Cost!K145"),
        ("AMD MI455X", "CMP_Supply_Amd", "COST_Pay_Amd", "Cost!K125", "Cost!K138", "Cost!K148"),
        ("SpaceX／xAI Colossus", "CMP_Supply_Spacex", "COST_Pay_Spacex", "Cost!K126", "Cost!K137", "Cost!K147"),
        ("Nscale", "CMP_Supply_Nscale", "COST_Pay_Nscale", "Cost!K127", None, None),
        ("Lambda", "CMP_Supply_Lambda", "COST_Pay_Lambda", "Cost!K128", None, None),
        ("Volta Infra", "CMP_Supply_Volta", "COST_Pay_Volta", "Cost!K129", None, None),
        ("Akamai", "CMP_Supply_Akamai", "COST_Pay_Akamai", "Cost!K130", None, None),
    ]
    crow = []
    for lab, g, p, tot, pros, gap in contracts:
        gg = V.get(gap) if gap else 0
        warn = ' class="neg"' if gap and abs(gg) > V.get("INP_179") else ""
        crow.append(f"<tr><td>{lab}</td><td>{cell(V, g, 2027, 'f2')}</td><td>{cell(V, g, 2030, 'f2')}</td><td>{cell(V, p, 2028)}</td><td>{cell(V, p, 2030)}</td>"
                    f"<td>{cell(V, tot)}</td><td>{cell(V, pros) if pros else '—'}</td><td{warn}>{cell(V, gap, None, 'pct0') if gap else '—'}</td></tr>")
    crow.append(f"<tr><td>既有雲端（2025 前簽）</td><td>{cell(V, 'CMP_Supply_Legacy', 2027, 'f2')}</td><td>{cell(V, 'CMP_Supply_Legacy', 2030, 'f2')}</td><td>{cell(V, 'COST_Pay_Legacy', 2028)}</td><td>{cell(V, 'COST_Pay_Legacy', 2030)}</td><td>—</td><td>—</td><td>—</td></tr>")
    crow.append(f"<tr><td>自建（Fluidstack $50B；資本支出）</td><td>{cell(V, 'CMP_Supply_Owned', 2027, 'f2')}</td><td>{cell(V, 'CMP_Supply_Owned', 2030, 'f2')}</td><td>{cell(V, 'COST_OwnedCapex', 2028)}</td><td>{cell(V, 'COST_OwnedCapex', 2030)}</td><td>—</td><td>—</td><td>—</td></tr>")
    crow.append(f"<tr class=\"em\"><td>合計</td><td>{cell(V, 'CMP_SupplyGW', 2027, 'f2')}</td><td>{cell(V, 'CMP_SupplyGW', 2030, 'f2')}</td><td>{cell(V, 'COST_Compute', 2028)}</td><td>{cell(V, 'COST_Compute', 2030)}</td>"
                f"<td>{cell(V, 'Cost!K132')}</td><td>{cell(V, 'Cost!K139')}</td><td>{cell(V, 'Cost!K141', None, 'pct0')}</td></tr>")
    S.append(f"""<section id="s4"><h2><span class="no">④</span>算力：總需求 vs 已簽合約＋自建供給（GW，IT 電力）</h2>
<div class="legend"><span><i style="background:var(--inf)"></i>有效推論 GW</span><span><i style="background:var(--rd)"></i>研發 GW（殘差）</span><span><i style="background:var(--con)"></i>合約供給</span><span><i style="background:var(--own)"></i>自建供給</span><span>每年左＝需求、右＝供給</span></div>
<div class="cw">{chart_compute(V)}</div>
{year_table(V, [("有效推論 GW", "CMP_InfGW_Eff", "f2", ""), ("研發 GW（殘差）", "CMP_RDGW", "f2", ""), ("總需求 GW", "CMP_DemandGW", "f2", "em"),
                ("合約供給 GW", "CMP_SupplyContractGW", "f2", ""), ("自建供給 GW", "CMP_Supply_Owned", "f2", ""), ("供給 GW 合計", "CMP_SupplyGW", "f2", "em"),
                ("供給 VR 等值 GW（命題分母）", "CMP_Supply_VReq", "f2", ""), ("推論可用 GW（容量上限）", "CMP_InfAvailGW", "f2", "")], head="GW")}
<p><b>供給 2028 年後遞減</b>：模型只計已簽合約，SpaceX 月費 2029-05 到期、Fluidstack 自建 2028 年後不再投入，供給 GW 由 2028 年 {n("CMP_SupplyGW", 2028, "f1")} 降到 2030 年 {n("CMP_SupplyGW", Y, "f1")}；這是 2029–2030 轉正的主因之一。研發 GW 是殘差（供給扣閒置後減推論），2027 年 {n("CMP_RDGW", 2027, "f1")} GW，比 Tokenomics 研發錨點高兩個數量級（Tokenomics 缺口）。</p>
<h3>逐家合約與招股書對帳</h3>
<div class="tw"><table><thead><tr><th>合作方</th><th>GW 2027</th><th>GW 2030</th><th>實付 2028</th><th>實付 2030</th><th>模型合約總額</th><th>招股書</th><th>模型 ÷ 招股書 − 1</th></tr></thead><tbody>
{"".join(crow)}
</tbody></table></div>
<p class="note">基準以「GW × 每 GW 年合約價 {n("INP_114", None, "f0")} × 持有比」或揭露金額計（D10）；逐家差距超過 {n("INP_179", None, "pct0")} 標紅並觸發 Checks WARN（AMD 模型 {n("Cost!K125", None, "f1")} 對招股書 {n("Cost!K138", None, "f0")}，SpaceX 錨在月費 $45B 對招股書「最高」{n("Cost!K137", None, "f1")}）。合計只差 {n("Cost!K141", None, "pct0")} 是誤差互相抵銷。逐家錨定招股書的敏感度（S4）：2026 差額 {n("COST_PropGap_VR_Anchor", 2026, "sf1")}、2030 {n("COST_PropGap_VR_Anchor", Y, "sf1")}，仍不需外部資金（見 ⑤）。機房租約 {n("Compute!C170", None, "f2")} GW 只列對照、不計供給。</p>
{V.src("CMP_InfGW_Eff", "CMP_RDGW", "CMP_DemandGW", "CMP_SupplyContractGW", "CMP_Supply_Owned", "CMP_SupplyGW", "CMP_Supply_VReq", "CMP_InfAvailGW", "Cost!K121", "Cost!K133", "Cost!K141", "Cost!K143", "COST_PropGap_VR_Anchor", "Compute!C170", "INP_114", "INP_179")}
</section>""")

    # ── ⑤ 現金與融資 ──
    scen_rows = [("基準（已到位 $100B）", "FND_CashEnd", "FND_ExtNeedCum"), ("S1 ＋IPO 約 $100B（2026-11）", "Funding!F49", "Funding!F50"),
                 ("S2 ＋條件式（Amazon、Google、NVIDIA、AMD）", "Funding!F54", "Funding!F55"), ("S3 ＋IPO＋條件式", "Funding!F59", "Funding!F60"),
                 ("S4 逐家錨定招股書金額", "Funding!F64", "Funding!F65"), ("S5 加計機房租約租金", "Funding!F69", "Funding!F70"),
                 ("S6 D22 非算力成本補足（約 +$7B）", "Funding!F74", "Funding!F75"), ("S7 S4＋S5＋S6 合併", "Funding!F79", "Funding!F80")]
    sr = []
    for lab, cash, cum in scen_rows:
        sr.append(f'<tr class="{"em" if cash == "FND_CashEnd" else ""}"><td>{lab}</td>' + "".join(f"<td>{cell(V, cash, y, 'f1')}</td>" for y in YEARS[1:])
                  + f"<td>{cell(V, cum, Y, 'f1')}</td></tr>")
    rt = []
    for rid, lab in (("F88", "Amazon"), ("F89", "Google"), ("F90", "Broadcom（債務）"), ("F91", "Microsoft"), ("F92", "NVIDIA（間接）"), ("F93", "AMD"), ("F94", "合計")):
        rt.append(f'<tr class="{"em" if rid == "F94" else ""}"><td>{lab}</td>' + "".join(
            f"<td>{cell(V, f'Funding!{rid}:{c}', None, 'pct0' if c in 'HI' else 'f1')}</td>" for c in "DEFGHI") + "</tr>")
    S.append(f"""<section id="s5"><h2><span class="no">⑤</span>現金路徑與融資</h2>
<div class="legend"><span><i style="background:var(--rev)"></i>年底現金（基準）</span><span><i style="background:var(--cost)"></i>年底現金（S7：錨定招股書＋機房租金＋非算力補足）</span><span><i style="background:var(--gap);height:3px"></i>最低現金（次年非算力 6 個月）</span></div>
<div class="cw">{chart_cash(V)}</div>
{year_table(V, [("營收淨額", "Funding!F01", "f1", ""), ("算力成本（現金）", "Funding!F02", "f1", ""), ("　其中自建資本支出", "Funding!F03", "f1", ""),
                ("非算力成本（不含股權報酬）", "Funding!F04", "f1", ""), ("自由現金流（股權報酬加回）", "FND_FCF", "sf1", "em"),
                ("期初現金（①）", "FND_CashOpen", "f1", ""), ("已到位融資（②）", "FND_Committed", "f1", ""), ("當年外部資金需求（③）", "FND_ExtNeed", "f1", "em"),
                ("年底現金", "FND_CashEnd", "f1", "em"), ("最低現金", "FND_MinCash", "f1", ""), ("現金餘裕（年底 − 最低）", "FND_Headroom", "f1", ""),
                ("累計外部資金需求（命題 2）", "FND_ExtNeedCum", "f1", "em")])}
<h3>已到位融資（2026；來源順序 ②）</h3>
<div class="tw"><table><thead><tr><th>來源</th><th>$B</th></tr></thead><tbody>
<tr><td>Series G（2026-02；GIC、Coatue 領投；估值 $380B）</td><td>{n("Funding!F12", 2026, "f0")}</td></tr>
<tr><td>Amazon Series G 無投票權特別股（2026 Q2；Amazon 10-Q，Verified）</td><td>{n("Funding!F13", 2026, "f0")}</td></tr>
<tr><td>Series H（2026-05；含先前承諾的雲端業者投資 {n("Funding!F15", 2026, "f0")}）</td><td>{n("Funding!F14", 2026, "f0")}</td></tr>
<tr><td>Microsoft $5B（2025-11，已在 2025 年底現金內）、Google 2026-04 即時 $10B（視為 Series H 一部分）</td><td>{n("Funding!F16", 2026, "f0")}／{n("Funding!F17", 2026, "f0")}</td></tr>
<tr class="em"><td>合計</td><td>{n("Funding!F18", 2026, "f0")}</td></tr>
</tbody></table></div>
<h3>條件式、IPO 與敏感度情境（同一規則；年底現金與 2030 累計外部資金需求）</h3>
<div class="tw"><table><thead><tr><th>情境</th><th>2026</th><th>2027</th><th>2028</th><th>2029</th><th>2030</th><th>累計外部資金</th></tr></thead><tbody>
{"".join(sr)}
</tbody></table></div>
<p class="note">條件式（未到位）：Amazon {n("Funding!F22", 2027, "f0")}（依算力交付）、Google {n("Funding!F23", 2027, "f0")}（績效）、NVIDIA 最高 {n("Funding!F24", 2026, "f0")}、AMD 認股 {n("Funding!F25", 2027, "f0")}；IPO 約 {n("Funding!F21", 2026, "f0")}。全部不計入基準。<b>或有負債（只列示）</b>：循環信貸 {n("Funding!F81", None, "f0")}、Broadcom 租賃融資最高 {n("Funding!F82", None, "f0")}、表外 Google TPU 融資 SPV {n("Funding!F83", None, "f0")}（單一來源，未證實），合計 {n("FND_Contingent", None, "f0")}；合約 2030 年後剩餘承諾 {n("FND_CommitAfter2030", None, "f1")}（招股書稱約 {n("Funding!F86", None, "pct0")} 不可取消）。</p>
<h3>回流對照（D18：只對照、不沖銷）</h3>
<div class="tw"><table><thead><tr><th>投資人兼算力供應商</th><th>已到位投資</th><th>條件式（未到位）</th><th>合約總額（招股書）</th><th>模型 2026–30 付款</th><th>已到位 ÷ 付款</th><th>（到位＋條件）÷ 合約</th></tr></thead><tbody>
{"".join(rt)}
</tbody></table></div>
<p class="note">策略投資人已投入的錢只相當於 2026–2030 付給它們的算力款的 {n("Funding!F94:H", None, "pct0")}：算力錢主要不是由賣算力的一方回流，而是由 2026 年的財務投資人與 2029 年起的營收提供。</p>
{V.src("Funding!F01", "FND_FCF", "FND_CashOpen", "FND_Committed", "FND_ExtNeed", "FND_CashEnd", "FND_MinCash", "FND_Headroom", "FND_ExtNeedCum", "Funding!F12", "Funding!F21", "Funding!F49", "Funding!F79", "FND_Contingent", "FND_CommitAfter2030", "Funding!F88")}
</section>""")

    # ── ⑥ 關鍵驅動 ──
    rows = [("基準", "base"),
            ("1 API 任務數成長 2027–30 取低", "api_task_lo"), ("1 API 任務數成長 2027–30 取高", "api_task_hi"),
            ("1 API 每任務 token 成長 取低", "tok_lo"), ("1 API 每任務 token 成長 取高", "tok_hi"),
            ("2 價格彈性（任務數）−1.05", "elast_strong"), ("2 價格彈性（任務數）−0.35", "elast_weak"), ("2 API 牌價年降 35%（觸發容量上限）", "price_cut"),
            ("3 每 GW 年合約價 8", "price_lo"), ("3 每 GW 年合約價 20", "price_hi"), ("3 爬坡 1 年", "ramp_lo"), ("3 爬坡 3 年", "ramp_hi"),
            ("3 AMD 合約只 3 年", "amd_short"), ("3 NonNV 產出比 TK 低", "nnv_lo"), ("3 NonNV 產出比 TK 高", "nnv_hi")]
    combos = [("任務數低＋彈性弱", "combo1"), ("任務數低＋合約價 20", "combo2"), ("彈性弱＋合約價 20＋爬坡 1 年", "combo3"),
              ("任務數低＋彈性弱＋合約價 20＋爬坡 1 年", "combo4"), ("上列＋每任務 token 低＋人數成長高＋抽成 30%", "combo5")]

    def srow(lab, sc, em=False):
        return (f'<tr class="{"em" if em else ""}"><td>{lab}</td><td>{n("REV_NetCapped", Y, "f1", sc)}</td><td>{n("COST_PropGap_VR", 2027, "sf1", sc)}</td>'
                f'<td>{n("COST_PropGap_VR", Y, "sf1", sc)}</td><td>{n("COST_Coverage", Y, "pct0", sc)}</td><td>{n("FND_ExtNeedCum", Y, "f1", sc)}</td>'
                f'<td>{n("FND_FirstGapYear", None, "yt", sc)}</td><td>{n("FND_CashNoExtTrough", None, "f1", sc)}</td></tr>')
    head = "<tr><th>情境</th><th>2030 營收淨額</th><th>2027 差額</th><th>2030 差額</th><th>2030 覆蓋率</th><th>累計外部資金</th><th>首次缺口年</th><th>現金谷底</th></tr>"
    S.append(f"""<section id="s6"><h2><span class="no">⑥</span>三個關鍵驅動、敏感度與「要成立必須…」</h2>
<p class="note">1＝API 需求成長（任務數、每任務 token；營收面最大驅動）；2＝價格反應（彈性、牌價年降；營收面＋容量上限）；3＝算力合約條款與產能（合約價、爬坡、AMD 期間、TPU／Trainium／AMD 產出比；成本與分母）。設定取自 Inputs 低／高欄與 Tokenomics 快照區間；全部由 Excel 重算。<b>任何單一驅動在區間內都不會產生外部資金需求。</b>差額單位 $B/GW/年（每 VR 等值 GW）。</p>
<div class="tw"><table><thead>{head}</thead><tbody>
{"".join(srow(lab, sc, sc == "base") for lab, sc in rows)}
</tbody></table></div>
<h3>能翻轉結論的組合</h3>
<div class="tw"><table><thead>{head}</thead><tbody>
{"".join(srow(lab, sc) for lab, sc in combos)}
</tbody></table></div>
<div class="must"><b>「不需外部資金」要成立，必須：</b>
<ol><li><b>2029 年後新增算力的支出不超過餘裕</b>（最大的前提；模型只計已簽合約，未建「新簽合約」機制）。若營收不變，2029 年新增算力支出超過 {n("COST_GapCash", 2029, "f1")}、2030 年超過 {n("COST_GapCash", Y, "f1")}（＝當年差額），差額就轉負；新增支出累計超過 {n("FND_Headroom", Y, "f1")}（2030 年現金餘裕；且各年不超過當年餘裕：2028 年 {n("FND_Headroom", 2028, "f1")}、2029 年 {n("FND_Headroom", 2029, "f1")}），才需要外部資金。對照：供給 GW 2028 → 2030 由 {n("CMP_SupplyGW", 2028, "f1")} 降到 {n("CMP_SupplyGW", Y, "f1")}，營收淨額卻由 {n("REV_NetCapped", 2028, "f1")} 增到 {n("REV_NetCapped", Y, "f1")}。</li>
<li><b>API 需求成長與價格彈性不同時落在下緣</b>：兩者同時取下緣 → 2030 差額 {n("COST_PropGap_VR", Y, "sf1", "combo1")}、累計外部資金需求 {n("FND_ExtNeedCum", Y, "f1", "combo1")}；再加合約價 20、爬坡 1 年 → {n("FND_ExtNeedCum", Y, "f1", "combo4")}（首次缺口 {n("FND_FirstGapYear", None, "yt", "combo4")}）。</li>
<li><b>2026 年的 ${n("FND_Committed", 2026, "f0")}B 已交割</b>（Series G／H 為公司新聞稿與律所公告，多個獨立來源）。</li></ol></div>
{V.src("REV_NetCapped", "COST_PropGap_VR", "COST_Coverage", "FND_ExtNeedCum", "FND_FirstGapYear", "FND_CashNoExtTrough", "COST_GapCash", "FND_Headroom", "CMP_SupplyGW")}
</section>""")

    # ── ⑦ 與 OpenAI v0.6 並排 ──
    pair = [("每 VR 等值 GW 營收淨額", "COST_PropRev_VR", "OAI_COST_PropRev_VR", "f1"), ("每 VR 等值 GW 全成本", "COST_PropFull_VR", "OAI_COST_PropFull_VR", "f1"),
            ("每 VR 等值 GW 差額（命題 1）", "COST_PropGap_VR", "OAI_COST_PropGap_VR", "sf1"), ("覆蓋率", "COST_Coverage", "OAI_COST_Coverage", "pct0"),
            ("累計外部資金需求（命題 2）", "FND_ExtNeedCum", "OAI_FND_ExtNeedCum", "f1"), ("營收淨額（截頂後）", "REV_NetCapped", "OAI_REV_NetCapped", "f1"),
            ("算力成本（現金）", "COST_Compute", "OAI_COST_Compute", "f1"), ("供給 VR 等值 GW", "CMP_Supply_VReq", "OAI_CMP_Supply_VReq", "f2"),
            ("年底現金", "FND_CashEnd", "OAI_FND_CashEnd", "f1")]
    pr = []
    for lab, a, o_, f in pair:
        em = "em" if "命題" in lab else ""
        pr.append(f'<tr class="{em}"><td>{lab}</td>' + "".join(f"<td>{cell(V, a, y, f)}<span class=\"note\">／</span>{cell(V, o_, y, f)}</td>" for y in YEARS[1:]) + "</tr>")
    S.append(f"""<section id="s7"><h2><span class="no">⑦</span>與 OpenAI v0.6 並排（同一 Tokenomics 快照、同一合約價）</h2>
<div class="legend"><span><i style="background:var(--rev)"></i>Anthropic v{ver}</span><span><i style="background:var(--oai)"></i>OpenAI v0.6</span><span>每 VR 等值 GW 差額，$B/GW/年；2025 兩家都是分母效應，略</span></div>
<div class="cw">{chart_pair(V)}</div>
<div class="tw"><table><thead><tr><th>Anthropic／OpenAI</th>{"".join(f"<th>{y}</th>" for y in YEARS[1:])}</tr></thead><tbody>
{"".join(pr)}
</tbody></table></div>
<p class="note">差距約一半來自每 GW 營收（Anthropic 以企業與 API 為主、免費用戶負擔小，2030 年每 VR 等值 GW 營收 {n("COST_PropRev_VR", Y)} 對 {n("OAI_COST_PropRev_VR", Y)}），另一半來自算力承諾規模（2030 算力成本 {n("COST_Compute", Y)} 對 {n("OAI_COST_Compute", Y)}；Anthropic 已簽合約 2029 年起遞減）。若 Anthropic 2029 年起比照 OpenAI 持續擴張算力，正差額會縮小或消失。兩家共用合約價 12，對 VR 世代偏低（Anthropic 雲端毛利率 2030 年 {n("COST_CloudGMPct", Y, "pct0")}）。OpenAI 數字取自 OAI_Link 快照（<code>{html.escape(str(oai_file))}</code>，SHA-256 <code>{html.escape(str(oai_sha))[:8]}…</code>），只作並排，不被任何計算頁引用。</p>
{V.src("COST_PropRev_VR", "OAI_COST_PropRev_VR", "OAI_COST_PropFull_VR", "OAI_COST_PropGap_VR", "OAI_COST_Coverage", "OAI_FND_ExtNeedCum", "OAI_REV_NetCapped", "OAI_COST_Compute", "OAI_CMP_Supply_VReq", "OAI_FND_CashEnd", "COST_CloudGMPct")}
</section>""")

    # ── ⑧ Andy 最該審的 5 項預設 ──
    S.append(f"""<section id="s8"><h2><span class="no">⑧</span>最該審的 5 項預設（取自 v{ver} 完成報告 ☆）</h2>
<ol>
<li><b>只計已簽合約（規格未列的慣例，與 OpenAI 同）</b>：使 2029–2030 轉正（供給 GW 2030 年 {n("CMP_SupplyGW", Y, "f1")} 低於 2028 年）。新增合約機制需 Andy 決定；上限見 ⑥ 的餘裕數字。</li>
<li><b>API 任務數成長（A2-13；與 OpenAI v0.6 同值）</b>：最大不確定；取低／高時 2030 差額 {n("COST_PropGap_VR", Y, "sf1", "api_task_lo")}／{n("COST_PropGap_VR", Y, "sf1", "api_task_hi")}、現金谷底 {n("FND_CashNoExtTrough", None, "f1", "api_task_lo")}／{n("FND_CashNoExtTrough", None, "f1", "api_task_hi")}。</li>
<li><b>2026 非算力成本可能低估（A4-17／D22）</b>：報導 2026 Q2 隱含低估 {n("COST_D22Under", None, "f1")}；補足後 2026 差額 {n("COST_PropGap_VR_D22", 2026, "sf1")}、現金谷底 {n("Funding!F74", 2028, "f1")}，結論不變。待公開版 S-1 校準。</li>
<li><b>每 GW 年合約價 12（D9）對 VR 世代偏低</b>：雲端毛利率 2027–2030 為 {n("COST_CloudGMPct", 2027, "pct0")} 至 {n("COST_CloudGMPct", Y, "pct0")}；取 20 時 2027 差額 {n("COST_PropGap_VR", 2027, "sf1", "price_hi")}、現金谷底 {n("FND_CashNoExtTrough", None, "f1", "price_hi")}。</li>
<li><b>2026 已到位 $100B 的組成與避免重複（A4-2～6）、逐家錨定招股書（A4-15）</b>：Series G 不計時谷底約少 {n("Funding!F12", 2026, "f0")}；逐家錨定（S4）谷底 {n("Funding!F64", 2028, "f1")}，仍無缺口。</li>
</ol>
<p class="note">A1–A4 共 75 項建置預設與規格 D1–D23 的全表：<code>{report}</code> 第 ④ 節、<code>docs/reports/{stamp}_v{ver}_對照.xlsx</code> 第 ④ 頁。</p>
{V.src("COST_D22Under", "COST_PropGap_VR_D22", "COST_CloudGMPct")}
</section>""")

    # ── ⑨ 資料與版本 ──
    S.append(f"""<section id="s9"><h2><span class="no">⑨</span>資料與版本</h2>
<dl>
<dt>Excel</dt><dd><code>{html.escape(model.name)}</code>（v{ver}；唯一計算引擎；本頁同資料夾附 xlsx）</dd>
<dt>Tokenomics</dt><dd>{html.escape(str(tk_ver))}（<code>{html.escape(str(tk_file))}</code>），master <code>{html.escape(str(tk_sha))}</code>，讀取日 {html.escape(str(tk_date))}</dd>
<dt>OpenAI 對照</dt><dd><code>{html.escape(str(oai_file))}</code>（v0.6），提交 <code>{html.escape(str(oai_commit))[:7]}</code></dd>
<dt>資料日期</dt><dd>SRC_ANT 共 {src_n} 列，最新文件日期 {src_date}；模型日期 {date_disp}</dd>
<dt>檢查</dt><dd>CHK_Errors＝{n("CHK_Errors", None, "f0")}、CHK_Warnings＝{n("CHK_Warnings", None, "f0")}（逐家合約對帳 AMD，預期）；Checks {chk_ids[0]}–{chk_ids[-1]}</dd>
<dt>期間與單位</dt><dd>FY2025–FY2030 曆年制，2025 為實際校準年；$B；GW＝IT 關鍵電力；VR 等值＝以 VR200 Sol 層級每 GW 產能換算</dd>
</dl>
<div class="must"><b>來源提醒</b>：2025–2026 的營收、算力支出、營業費用、現金與合約承諾多數取自<b>尚未公開的 IPO 招股書草稿</b>，經 Reuters、Fortune 等轉述（SRC_ANT 中 {src_pros} 列的來源或備註引用招股書／說明書），標記為 <span class="tag">Interested-party</span>（公司為上市募資有誘因呈現較佳數字）。公開版 S-1 上 EDGAR 後應改為 Verified 並覆核。</div>
<p class="note"><span class="tag">Verified</span>已查核原文 <span class="tag">Interested-party</span>公司或利害關係方說法 <span class="tag">Analogy</span>類比推估（附區間） <span class="tag">Assumed</span>假設（附區間） <span class="tag">Derived</span>由其他數字推得 <span class="tag">Decision</span>建模決定。公司原始數據在 SRC_ANT、算力物理取自 Tokenomics（TK_Link）、假設在 Inputs、OpenAI 對照在 OAI_Link。</p>
<p class="note">本頁所有數字皆讀自 LibreOffice 重算後的 Excel（具名範圍或頁＋列 ID），HTML 不含計算；敏感度情境定義見 <code>tools/html_scenarios.yaml</code>；完整報告 <code>{report}</code>。</p>
</section>""")

    body = "\n".join(S)
    return f"""<!DOCTYPE html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="color-scheme" content="light dark">
<title>Anthropic 收支模型 v{ver}</title>
<style>{CSS}</style>
</head>
<body>
<main>
{body}
<footer>Anthropic 收支模型 v{ver}｜產生器 tools/build_html.py｜{date_disp}</footer>
</main>
</body>
</html>
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--outdir", type=Path, default=REPO / "dist")
    a = ap.parse_args()
    h, x, V = build(a.outdir)
    print(f"HTML：{h}（{h.stat().st_size:,} bytes；數字 {len(V.used)} 個）")
    print(f"Excel：{x}")


if __name__ == "__main__":
    main()
