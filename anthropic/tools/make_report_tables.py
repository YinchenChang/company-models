#!/usr/bin/env python3
"""產生每段報告的對照 Excel（CLAUDE.md 第 5 節）與報告用的 Markdown 表（Anthropic v0.1）。

用法：python3 tools/make_report_tables.py --stage A2|A3|A4 --out docs/reports/20261008_v0.1-An_對照.xlsx --report <報告.md> [--md build_out/report_tables.md]
三頁：①本段新增或變動的每一列（頁、編號／ID、列名、值、來源／標記）；②本段關鍵輸出 FY2025–2030；③已套用的預設。
A4（v0.1 完成報告）另加：④A1–A4 已套用的預設彙總（報告中以「| # | 段 |」開頭的表；規格層級另一表以「| ID | 規格預設」開頭）；⑤關鍵驅動敏感度（build_out/sensitivity.json，tools/sensitivity.py 產生）。
數值一律讀現行活頁簿（LibreOffice 重算後的快取值）；③的文字取自 docs/reports/<報告>.md 的「已套用的預設」表（同一份內容，不另抄）。
"""
import argparse
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
from engine import current_model_path  # noqa: E402

BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="FFEDEDED")
YEARS = ["2025", "2026", "2027", "2028", "2029", "2030"]
# ② 關鍵輸出：（頁, 列鍵名稱或具名範圍, 顯示名稱）
KEY_OUT = [
    ("Revenue", "REV_Sub_Pro", "訂閱：Pro"), ("Revenue", "REV_Sub_Max5", "訂閱：Max 5x"), ("Revenue", "REV_Sub_Max20", "訂閱：Max 20x"),
    ("Revenue", "REV_SubConsumer", "個人訂閱合計"), ("Revenue", "REV_Sub_TeamStd", "訂閱：Team 標準"), ("Revenue", "REV_Sub_TeamPrem", "訂閱：Team Premium"),
    ("Revenue", "REV_Sub_Ent", "訂閱：Enterprise 席位"), ("Revenue", "REV_SubSeats", "企業席位訂閱合計"), ("Revenue", "REV_Sub", "訂閱合計"),
    ("Revenue", "REV_API_Top", "API：頂層（Opus）"), ("Revenue", "REV_API_Mid", "API：中層（Sonnet）"), ("Revenue", "REV_API_Low", "API：低層（Haiku）"),
    ("Revenue", "REV_API", "API 合計"), ("Revenue", "REV_Other", "其他"), ("Revenue", "REV_Ads", "廣告（D3＝0）"),
    ("Revenue", "REV_Gross", "營收總額（報導口徑）"), ("Revenue", "REV_Channel", "雲端通路營收"), ("Revenue", "REV_PartnerShare", "雲端平台抽成"),
    ("Revenue", "REV_Net", "營收淨額"), ("Revenue", "REV_Consumer", "個人（Pro、Max）"), ("Revenue", "REV_Enterprise", "企業（席位＋API）"),
    ("Revenue", "REV_NetCapped", "營收淨額（截頂後；A2 佔位＝1）"), ("Revenue", "REV_ApiPrice", "API 組合有效單價（$/M）"),
    ("Demand", "DEM_Users_PaidConsumer", "個人付費人數（M）"), ("Demand", "DEM_Users_Free", "Free 人數（M）"), ("Demand", "DEM_Seats_Total", "企業席位數（M）"),
    ("Demand", "DEM_Tok_Paid_Top", "付費 token：頂層（T）"), ("Demand", "DEM_Tok_Paid_Mid", "付費 token：中層（T）"), ("Demand", "DEM_Tok_Paid_Low", "付費 token：低層（T）"),
    ("Demand", "DEM_Tok_Free", "免費 token（T）"), ("Demand", "DEM_Tok_Paid", "付費 token（T）"), ("Demand", "DEM_Tok_Total", "token 合計（T）"),
    ("Demand", "DEM_Tok_API", "其中 API 計費 token（T）"),
    ("OAI_Link", "OAI_REV_Gross", "OpenAI v0.6 營收總額（對照）"), ("OAI_Link", "OAI_REV_Net", "OpenAI v0.6 營收淨額（對照）"),
    ("OAI_Link", "OAI_DEM_Tok_Total", "OpenAI v0.6 token 合計（對照）"),
]

# A3：② 第一部分＝命題表（每 VR 等值 GW），其後為算力與成本關鍵輸出，最後為 OpenAI v0.6 並排
KEY_OUT_A3 = [
    ("Cost", "COST_PropRev_VR", "命題表｜每 VR 等值 GW：營收淨額（$B/GW/年）"), ("Cost", "COST_PropCompute_VR", "命題表｜每 VR 等值 GW：算力成本（現金）"),
    ("Cost", "COST_PropNonComp_VR", "命題表｜每 VR 等值 GW：非算力成本（不含股權報酬）"), ("Cost", "COST_PropSBC_VR", "命題表｜每 VR 等值 GW：股權報酬"),
    ("Cost", "COST_PropFull_VR", "命題表｜每 VR 等值 GW：全成本（含股權報酬）"), ("Cost", "COST_PropFullExSBC_VR", "命題表｜每 VR 等值 GW：全成本（不含股權報酬）"),
    ("Cost", "COST_PropGap_VR", "命題表｜每 VR 等值 GW：差額（含股權報酬；命題 1）"), ("Cost", "COST_PropGapExSBC_VR", "命題表｜每 VR 等值 GW：差額（不含股權報酬）"),
    ("Cost", "COST_Coverage", "命題表｜覆蓋率（營收淨額 ÷ 全成本，含股權報酬）"), ("Cost", "COST_CoverageExSBC", "命題表｜覆蓋率（不含股權報酬）"),
    ("Cost", "COST_PropGapEcon_VR", "命題表（經濟口徑）｜每 VR 等值 GW 差額（含股權報酬）"),
    ("Compute", "CMP_SupplyGW", "供給 GW（實體）"), ("Compute", "CMP_Supply_VReq", "供給 VR 等值 GW（命題分母）"), ("Compute", "CMP_VReqFactor", "機隊 VR 等值係數"),
    ("Compute", "CMP_Eta", "η"), ("Compute", "CMP_TokGW", "推論 GW（token 換算）"), ("Compute", "CMP_InfGW_Eff", "有效推論 GW（未截頂）"),
    ("Compute", "CMP_InfGW", "推論 GW（截頂後）"), ("Compute", "CMP_RDGW", "研發 GW"), ("Compute", "CMP_InfAvailGW", "推論可用 GW"),
    ("Compute", "CMP_CapFactor", "容量上限係數"), ("Compute", "CMP_Share_TPUv7", "組合：TPU v7／次世代"), ("Compute", "CMP_Share_Hopper", "組合：Hopper"),
    ("Compute", "CMP_Share_MI455X", "組合：AMD MI455X"),
    ("Cost", "COST_ContractPay", "合約實付合計（$B）"), ("Cost", "COST_OwnedCapex", "自建資本支出（Fluidstack）"), ("Cost", "COST_Compute", "算力成本（現金）"),
    ("Cost", "COST_ComputeEcon", "算力成本（經濟口徑）"), ("Cost", "COST_SupplierHold", "供應商持有成本（TK）"), ("Cost", "COST_CloudGM", "雲端毛利"),
    ("Cost", "COST_Headcount", "員工人數（年均）"), ("Cost", "COST_NonCompExSBC", "非算力成本（不含股權報酬）"), ("Cost", "COST_SBC", "股權報酬"),
    ("Cost", "COST_FullCash", "全成本（含股權報酬）"), ("Revenue", "REV_NetCapped", "營收淨額（截頂後）"), ("Cost", "COST_GapCash", "差額（營收淨額 − 全成本）"),
    ("Cost", "COST_InfCompute", "推論算力成本"), ("Cost", "COST_RDCompute", "研發算力成本"),
    ("OAI_Link", "OAI_CMP_SupplyGW", "OpenAI v0.6｜供給 GW"), ("OAI_Link", "OAI_CMP_Supply_VReq", "OpenAI v0.6｜供給 VR 等值 GW"),
    ("OAI_Link", "OAI_COST_Compute", "OpenAI v0.6｜算力成本"), ("OAI_Link", "OAI_COST_FullCash", "OpenAI v0.6｜全成本"),
    ("OAI_Link", "OAI_COST_PropRev_VR", "OpenAI v0.6｜每 VR 等值 GW 營收淨額"), ("OAI_Link", "OAI_COST_PropFull_VR", "OpenAI v0.6｜每 VR 等值 GW 全成本"),
    ("OAI_Link", "OAI_COST_PropGap_VR", "OpenAI v0.6｜每 VR 等值 GW 差額"), ("OAI_Link", "OAI_COST_Coverage", "OpenAI v0.6｜覆蓋率"),
]


# A4：② 第一部分＝融資（命題 2），其後為命題 1 三種口徑、情境、反向、OpenAI 並排
KEY_OUT_A4 = [
    ("Funding", "FND_FCF", "自由現金流（營收淨額 − 算力 − 非算力，不含股權報酬）"), ("Funding", "FND_Committed", "已到位融資（來源順序 ②）"),
    ("Funding", "FND_EquityIn", "計入基準的股權流入"), ("Funding", "FND_MinCash", "最低現金（D17）"), ("Funding", "FND_CashOpen", "期初現金"),
    ("Funding", "FND_ExtNeed", "當年外部資金需求（命題 2）"), ("Funding", "FND_ExtNeedCum", "累計外部資金需求（命題 2）"), ("Funding", "FND_CashEnd", "年底現金"),
    ("Funding", "FND_Headroom", "現金餘裕（年底現金 − 最低現金）"), ("Funding", "FND_CashNoExt", "只靠已到位融資的年底現金"),
    ("Cost", "COST_PropGap_VR", "命題 1｜每 VR 等值 GW 差額（現金，含股權報酬）"), ("Cost", "COST_PropGapEcon_VR", "命題 1｜經濟口徑每 VR 等值 GW 差額"),
    ("Cost", "COST_PropGap_GW", "命題 1｜每實體 GW 差額"), ("Cost", "COST_Coverage", "覆蓋率"), ("Cost", "COST_GapCash", "差額（$B）"),
    ("Cost", "COST_ComputeAnchor", "敏感度｜算力成本（逐家錨定招股書）"), ("Cost", "COST_PropGap_VR_Anchor", "敏感度｜每 VR 等值 GW 差額（錨定招股書）"),
    ("Cost", "COST_LeaseRent", "敏感度｜機房租約租金"), ("Cost", "COST_PropGap_VR_Rent", "敏感度｜每 VR 等值 GW 差額（加計租金）"),
    ("Cost", "COST_D22DeltaNC", "敏感度｜D22 非算力成本增加"), ("Cost", "COST_PropGap_VR_D22", "敏感度｜每 VR 等值 GW 差額（D22 補足）"),
    ("Funding", "FND_ExtNeedCumIPO", "情境 S1｜＋IPO：累計外部資金需求"), ("Funding", "FND_CashEndIPO", "情境 S1｜＋IPO：年底現金"),
    ("Funding", "FND_ExtNeedCumCond", "情境 S2｜＋條件式：累計外部資金需求"), ("Funding", "FND_ExtNeedCumAll", "情境 S3｜＋IPO＋條件式：累計外部資金需求"),
    ("Funding", "FND_CashEndAll", "情境 S3｜年底現金"), ("Funding", "FND_ExtNeedCumAnchor", "情境 S4｜錨定招股書：累計外部資金需求"),
    ("Funding", "FND_ExtNeedCumRent", "情境 S5｜機房租金：累計外部資金需求"), ("Funding", "FND_ExtNeedCumD22", "情境 S6｜D22：累計外部資金需求"),
    ("Funding", "FND_ExtNeedCumAdverse", "情境 S7｜S4＋S5＋S6：累計外部資金需求"),
    ("Reverse", "RVS_Target", "反向｜管理層營收目標"), ("Reverse", "RVS_Gap", "反向｜目標 − 正向"), ("Reverse", "RVS_MultAPI", "反向｜只靠 API 倍數"),
    ("Reverse", "RVS_MultSub", "反向｜只靠訂閱倍數"), ("Reverse", "RVS_MultProp", "反向｜兩線等比例倍數"), ("Reverse", "RVS_FCF", "反向｜自由現金流"),
    ("Reverse", "RVS_ExtNeedCum", "反向｜累計外部資金需求"), ("Reverse", "RVS_CashEnd", "反向｜年底現金"),
    ("OAI_Link", "OAI_COST_PropGap_VR", "OpenAI v0.6｜每 VR 等值 GW 差額"), ("OAI_Link", "OAI_COST_Coverage", "OpenAI v0.6｜覆蓋率"),
    ("OAI_Link", "OAI_FND_ExtNeed", "OpenAI v0.6｜當年外部資金需求"), ("OAI_Link", "OAI_FND_ExtNeedCum", "OpenAI v0.6｜累計外部資金需求"),
    ("OAI_Link", "OAI_FND_CashEnd", "OpenAI v0.6｜年底現金"),
]


def table_from_md(md: Path, head: str):
    """讀報告中第一個以 head 開頭的 Markdown 表。"""
    out, on = [], False
    for ln in md.read_text(encoding="utf-8").splitlines():
        if ln.startswith(head):
            on = True
            out.append([c.strip() for c in ln.strip("|").split("|")])
            continue
        if on:
            if not ln.startswith("|"):
                break
            if re.match(r"^\|[-| ]+\|$", ln):
                continue
            out.append([c.strip() for c in ln.strip("|").split("|")])
    return out


def name_values(wb, names, n):
    sh, rng = names[n].split("!")
    return [c.value for row in wb[sh][rng.replace("$", "")] for c in row]


def defaults_from_md(md: Path):
    """讀報告的「已套用的預設」表（第一個以「| # | 問題」開頭的表）。"""
    lines = md.read_text(encoding="utf-8").splitlines()
    out, on = [], False
    for ln in lines:
        if ln.startswith("| # | 問題"):
            on = True
            out.append([c.strip() for c in ln.strip("|").split("|")])
            continue
        if on:
            if not ln.startswith("|"):
                break
            if re.match(r"^\|[-| ]+\|$", ln):
                continue
            out.append([c.strip() for c in ln.strip("|").split("|")])
    return out


def sheet1(a, wb):
    names = {k: v.attr_text for k, v in wb.defined_names.items()}
    out = openpyxl.Workbook()

    # ① 本段新增或變動的每一列
    ws = out.active
    ws.title = "①新增與變動列"
    ws.append(["頁", "編號／ID", "列名", "單位", "值（2025 或單值）", "值（2030）", "來源／標記", "說明"])
    if a.stage == "A4":
        stage_a4_rows(wb, ws)
    elif a.stage == "A3":
        stage_a3_rows(wb, ws)
    else:
        stage_a2_rows(wb, ws)
    for c, w in zip("ABCDEFGH", (10, 22, 60, 14, 16, 14, 18, 70)):
        ws.column_dimensions[c].width = w
    key_out = {"A3": KEY_OUT_A3, "A4": KEY_OUT_A4}.get(a.stage, KEY_OUT)
    return out, names, key_out


def stage_a3_rows(wb, ws):
    """A3 新增或變動：Inputs INP_109 起（A3 新增）、TK_PUE、Revenue 容量上限列、Compute、Cost 每列、Checks C49 起。"""
    s = wb["Inputs"]
    for r in range(5, s.max_row + 1):
        iid = s.cell(r, 1).value
        if iid and int(iid[4:]) >= 109:
            ws.append(["Inputs", iid, f"{s.cell(r, 2).value}（{s.cell(r, 3).value}）", s.cell(r, 4).value, s.cell(r, 5).value, None,
                       s.cell(r, 8).value, f"低 {s.cell(r, 6).value}／高 {s.cell(r, 7).value}；{s.cell(r, 9).value}"])
    ws.append(["Inputs", "INP_016", "容量上限係數（A2 佔位）", "倍", None, None, "退役", "A3：REV_CapFactor 改接 CMP_CapFactor；號碼不重用"])
    s = wb["TK_Link"]
    for r in range(10, s.max_row + 1):
        if str(s.cell(r, 2).value or "").startswith("TK_PUE"):
            ws.append(["TK_Link", s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 10).value, None, s.cell(r, 6).value, s.cell(r, 1).value])
    s = wb["Revenue"]
    for r in range(7, s.max_row + 1):
        if s.cell(r, 2).value and str(s.cell(r, 2).value).startswith("容量上限係數"):
            ws.append(["Revenue", s.cell(r, 1).value, s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 9).value, "公式（變動）", "＝Compute CMP_CapFactor"])
    for sh, pre in (("Compute", "C"), ("Cost", "K")):
        s = wb[sh]
        for r in range(7, s.max_row + 1):
            code, lab = s.cell(r, 1).value, s.cell(r, 2).value
            if code and lab and re.fullmatch(pre + r"\d+", str(code)):
                ws.append([sh, code, lab, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 9).value, "公式／參數", s.cell(r, 11).value])
    s = wb["Checks"]
    for r in range(5, s.max_row + 1):
        cid = str(s.cell(r, 1).value or "")
        if re.fullmatch(r"C\d+", cid) and int(cid[1:]) >= 49:
            ws.append(["Checks", cid, s.cell(r, 2).value, None, s.cell(r, 3).value, None, s.cell(r, 5).value, s.cell(r, 6).value])


def stage_a4_rows(wb, ws):
    """A4 新增或變動：Inputs INP_165 起、Cost 第十二節起、Funding、Reverse 每列、Checks C83 起。"""
    s = wb["Inputs"]
    for r in range(5, s.max_row + 1):
        iid = s.cell(r, 1).value
        if iid and int(iid[4:]) >= 165:
            ws.append(["Inputs", iid, f"{s.cell(r, 2).value}（{s.cell(r, 3).value}）", s.cell(r, 4).value, s.cell(r, 5).value, None,
                       s.cell(r, 8).value, f"低 {s.cell(r, 6).value}／高 {s.cell(r, 7).value}；{s.cell(r, 9).value}"])
    s = wb["Cost"]
    on = False
    for r in range(7, s.max_row + 1):
        code, lab = s.cell(r, 1).value, s.cell(r, 2).value
        if isinstance(code, str) and code.startswith("十二"):
            on = True
        if lab == "對照：每實體 GW 差額（含股權報酬）":
            ws.append(["Cost", code, lab, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 9).value, "公式（變動：新增具名範圍 COST_PropGap_GW）", s.cell(r, 11).value])
        if on and code and lab and re.fullmatch(r"K\d+", str(code)):
            ws.append(["Cost", code, lab, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 9).value, "公式（A4 新增）", s.cell(r, 11).value])
    for sh, pre in (("Funding", "F"), ("Reverse", "X")):
        s = wb[sh]
        for r in range(7, s.max_row + 1):
            code, lab = s.cell(r, 1).value, s.cell(r, 2).value
            if code and lab and re.fullmatch(pre + r"\d+", str(code)):
                ws.append([sh, code, lab, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 9).value, "公式", s.cell(r, 11).value])
    s = wb["Checks"]
    for r in range(5, s.max_row + 1):
        cid = str(s.cell(r, 1).value or "")
        if re.fullmatch(r"C\d+", cid) and int(cid[1:]) >= 83:
            ws.append(["Checks", cid, s.cell(r, 2).value, None, s.cell(r, 3).value, None, s.cell(r, 5).value, s.cell(r, 6).value])


def stage_a2_rows(wb, ws):
    s = wb["SRC_ANT"]
    for r in range(5, s.max_row + 1):
        sid, status = s.cell(r, 1).value, s.cell(r, 22).value
        if not sid:
            continue
        if int(sid[8:]) >= 377 or status != "觀測":
            ws.append(["SRC_ANT", sid, s.cell(r, 4).value, s.cell(r, 8).value, s.cell(r, 5).value, None, s.cell(r, 16).value,
                       ("A2 新增（只增不重用）" if int(sid[8:]) >= 377 else status)])
    s = wb["Inputs"]
    for r in range(5, s.max_row + 1):
        if s.cell(r, 1).value:
            ws.append(["Inputs", s.cell(r, 1).value, f"{s.cell(r, 2).value}（{s.cell(r, 3).value}）", s.cell(r, 4).value, s.cell(r, 5).value, None,
                       s.cell(r, 8).value, f"低 {s.cell(r, 6).value}／高 {s.cell(r, 7).value}；{s.cell(r, 9).value}"])
    s = wb["TK_Link"]
    for r in range(10, s.max_row + 1):
        if s.cell(r, 2).value and (str(s.cell(r, 2).value).startswith("TK_NNV_") or s.cell(r, 1).value in ("IF_HdrTask", "IF_TaskLen", "IF_TaskTokFresh", "IF_TaskTokCached", "IF_TaskTokDec")):
            vals = [s.cell(r, 10 + k).value for k in range(int(s.cell(r, 5).value or 0))]
            ws.append(["TK_Link", s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value, ", ".join(map(str, vals)), None, s.cell(r, 6).value, s.cell(r, 1).value])
    s = wb["OAI_Link"]
    for r in range(10, s.max_row + 1):
        if s.cell(r, 1).value:
            ws.append(["OAI_Link", s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 10).value, s.cell(r, 4).value, s.cell(r, 9).value, "OpenAI v0.6 快照", s.cell(r, 11).value])
    for sh in ("Demand", "Revenue"):
        s = wb[sh]
        for r in range(7, s.max_row + 1):
            code, lab = s.cell(r, 1).value, s.cell(r, 2).value
            if code and lab and re.fullmatch(r"[DR]\d\d", str(code)):
                ws.append([sh, code, lab, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 9).value, "公式", s.cell(r, 11).value])
    s = wb["Checks"]
    for r in range(5, s.max_row + 1):
        if s.cell(r, 1).value and str(s.cell(r, 1).value).startswith("C"):
            ws.append(["Checks", s.cell(r, 1).value, s.cell(r, 2).value, None, s.cell(r, 3).value, None, s.cell(r, 5).value, s.cell(r, 6).value])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--stage", default="A2")
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--report", required=True, type=Path, help="報告 .md（讀其中的「已套用的預設」表）")
    ap.add_argument("--md", type=Path, help="輸出 ② 的 Markdown 表（貼入報告）")
    a = ap.parse_args()
    model = current_model_path()
    wb = openpyxl.load_workbook(model, data_only=True)
    out, names, key_out = sheet1(a, wb)

    # ② 關鍵輸出
    ws2 = out.create_sheet("②關鍵輸出 FY2025–2030")
    ws2.append(["項目", "具名範圍"] + YEARS)
    md = ["| 項目 | " + " | ".join(YEARS) + " |", "|---|" + "---|" * 6]
    for sh, n, lab in key_out:
        v = name_values(wb, names, n)
        ws2.append([lab, n] + v)
        fmt = (lambda x: f"{x:,.0f}") if "（T）" in lab else (lambda x: f"{x:,.2f}") if "（M）" in lab or "$/M" in lab else (lambda x: f"{x:,.2f}")
        md.append(f"| {lab} | " + " | ".join(fmt(x) if isinstance(x, (int, float)) else "" for x in v) + " |")
    ws2.column_dimensions["A"].width = 52
    ws2.column_dimensions["B"].width = 24

    # ③ 已套用的預設
    ws3 = out.create_sheet("③已套用的預設")
    for row in defaults_from_md(a.report):
        ws3.append(row)
    for c, w in zip("ABCDE", (5, 40, 60, 40, 60)):
        ws3.column_dimensions[c].width = w
    if a.stage == "A4":
        ws4 = out.create_sheet("④A1–A4 預設彙總")
        for row in table_from_md(a.report, "| ID | 規格預設"):
            ws4.append(row)
        ws4.append([])
        for row in table_from_md(a.report, "| # | 段 |"):
            ws4.append(row)
        for c, w in zip("ABCDEF", (6, 8, 40, 60, 40, 60)):
            ws4.column_dimensions[c].width = w
        import json
        sj = json.loads((REPO / "build_out" / "sensitivity.json").read_text(encoding="utf-8"))
        ws5 = out.create_sheet("⑤敏感度")
        ws5.append(["排名", "驅動", "群組", "設定（低）", "設定（高）", "2030 營收淨額 低", "高", "2027 差額 低", "高", "2030 差額 低", "高",
                    "累計外部資金需求 低", "高", "現金谷底 低", "高"])
        ws5.append([0, "基準", "", "", "", sj["base"]["rev30"], None, sj["base"]["gap27"], None, sj["base"]["gap30"], None, sj["base"]["cum30"], None,
                    sj["base"]["trough"], None])
        for i, d in enumerate(sj["drivers"], 1):
            ws5.append([i, d["driver"], d["group"], str(d["lo_set"]), str(d["hi_set"]), d["lo"]["rev30"], d["hi"]["rev30"], d["lo"]["gap27"], d["hi"]["gap27"],
                        d["lo"]["gap30"], d["hi"]["gap30"], d["lo"]["cum30"], d["hi"]["cum30"], d["lo"]["trough"], d["hi"]["trough"]])
        ws5.column_dimensions["B"].width = 36
    for w_ in out.worksheets:
        for c in w_[1]:
            c.font, c.fill = BOLD, HDR
        for row in w_.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        w_.freeze_panes = "A2"
    a.out.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.out)
    if a.md:
        a.md.parent.mkdir(parents=True, exist_ok=True)
        a.md.write_text("\n".join(md) + "\n", encoding="utf-8")
    print("saved", a.out, {w_.title: w_.max_row - 1 for w_ in out.worksheets})


if __name__ == "__main__":
    main()
