#!/usr/bin/env python3
"""v0.1.4 對照 Excel（CLAUDE.md 第 5 節）：①TK 39 個 DIFF 名稱新舊值與用途 ②v0.1.3 → v0.1.4 差異拆解 ③η 診斷補充前後
④逐格比對（範圍外變動）⑤已套用的預設。

用法：python3 tools/make_report_v014.py --new model/20261022_Anthropic_v0.1.4.xlsx --old model/archive/20261022_Anthropic_v0.1.3.xlsx \
        --tkdiff /tmp/work/tkdiff.json --out docs/reports/20261022_v0.1.4_對照.xlsx [--json /tmp/work/report_v014.json]
數值一律讀 LibreOffice 重算後的快取值；本程式不做模型計算。逐群拆解以 tools/html_scenarios.yaml 的情境（v013、v14_only_luna、v14_only_rest）
由 Excel 重算取得（TK 值取自上一版 Excel，不在本程式硬編碼）。
"""
import argparse
import json
import re
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "tools"))
sys.path.insert(0, str(REPO / "tests" / "parity"))
from build_html import Locator, load_scenarios, resolve_overrides  # noqa: E402

YEARS = ["2025", "2026", "2027", "2028", "2029", "2030"]
KEY_OUT = [
    ("COST_PropGap_VR", "每 VR 等值 GW 差額（含股權報酬，$B/GW/年）"), ("COST_Coverage", "覆蓋率（營收淨額 ÷ 全成本）"),
    ("CMP_Supply_VReq", "供給 VR 等值 GW"), ("CMP_SupplyGW", "供給 GW"), ("CMP_Eta", "η"), ("CMP_InfGW_Eff", "有效推論 GW"),
    ("COST_Compute", "算力成本（現金，$B）"), ("REV_NetCapped", "營收淨額（截頂後，$B）"), ("FND_FCF", "營運自由現金流（$B）"),
    ("FND_CashEnd", "年底現金（基準，$B）"), ("FND_ExtNeedCum", "累計外部資金需求（基準，$B）"),
]
SCALAR_OUT = [
    ("FND_L1_ExtTotal", "資金第 1 層累計外部資金需求（$B）"), ("FND_L2_ExtTotal", "資金第 2 層累計外部資金需求（$B；＝基準）"), ("FND_L3_ExtTotal", "資金第 3 層累計外部資金需求（$B）"),
    ("FND_L1_Trough", "第 1 層現金谷底（$B）"), ("FND_L2_Trough", "第 2 層現金谷底（$B）"), ("FND_L3_Trough", "第 3 層現金谷底（$B）"),
    ("FND_L1_GapYear", "第 1 層最低現金缺口首次年"), ("FND_L2_GapYear", "第 2 層最低現金缺口首次年"), ("FND_L3_GapYear", "第 3 層最低現金缺口首次年"),
    ("FND_FCFFirstPos", "營運自由現金流首次轉正年"), ("FND_CashNoExtTrough", "只靠已到位融資的現金谷底（$B）"),
]
EXP_OUT = [(f"COST_PropGap_VR_Exp{k}", f"持續擴張・{lab}：每 VR 等值 GW 差額") for k, lab in (("None", "無需求"), ("Half", "部分變現"), ("Full", "完全變現"))] + \
          [(f"FND_L2_ExtTotal_Exp{k}", f"持續擴張・{lab}：第 2 層累計外部資金需求（$B）") for k, lab in (("None", "無需求"), ("Half", "部分變現"), ("Full", "完全變現"))]
ETA_ROWS = [("ETA_Eta", "η（2025）"), ("ETA_fP", "f_P 計價口徑"), ("ETA_fU", "f_U 利用率（不計生產折減）"), ("ETA_fW", "f_W 工作負載組合"),
            ("ETA_Resid", "未解釋殘差（不計生產折減；v0.1.3 口徑）"), ("ETA_fU_Prod", "f_U' 利用率（計入生產折減）"), ("ETA_Resid2", "未解釋殘差（計入生產折減；v0.1.4）")]
DEFAULTS = [
    ("η 拆解是否加「TK 絕對產能水準」因子", "不加（Andy 2026-10-22 照預設）", "加第四因子", "不加則 TK 整體產能改變時全數落入殘差；v19.1 時殘差 ×0.99（見 ③ 版本小表）"),
    ("f_U 是否計入 TK 生產折減", "計入（f_U' ＝ f_U × IF_ProdDerate 0.85；Andy 2026-10-22 照預設）", "不計（v0.1.3 口徑，EtaDiag 第五節原列保留）", "殘差 ×1.176（0.440 → 0.517）；η 不變"),
    ("TK 生產折減的讀取", "TK 具名範圍 IF_ProdDerate（Interface C355；TK_Link 尾端新增）", "讀表（非具名）", "有具名範圍，無 Tokenomics 缺口"),
    ("殘差對照 TK 的方式", "本模型殘差倒數（TK 方向）÷ TK L1_GapUnexplained，基準對基準；區間端點不相除", "區間端點相除", "端點相除會把兩邊區間疊成更寬的範圍，無資訊"),
    ("TK 快照取用", "master c2ecb71 的 CURRENT v19.1（工作單指定；與 master 現況相同）", "—", "—"),
    ("OpenAI 並排快照", "維持 v0.6（取 archive/20261008_OpenAI_v0.6.xlsx）", "改 v0.6.1／v0.6.2", "只影響並排列"),
    ("dist/_待刪除/", "git rm（Drive 仍有副本；Andy 2026-10-22 照預設）", "保留", "無數值影響"),
    ("HTML 位置", "v0.1.4 摘要併入「①之三」前端並改名「v0.1.3＋v0.1.4」；①之二的 v0.1.2 欄以 tk19_back 保持當時數字", "另開 ①之四", "僅版面"),
]


def named(wb, name):
    dn = wb.defined_names[name].attr_text.replace("$", "")
    sh, rg = dn.split("!")
    c = wb[sh.strip("'")][rg]
    if not isinstance(c, tuple):
        c = ((c,),)
    return [x.value for r in c for x in r]


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def same(a, b):
    return a == b or (num(a) and num(b) and abs(a - b) <= 1e-9 * max(1, abs(a), abs(b)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--new", required=True, type=Path)
    ap.add_argument("--old", required=True, type=Path)
    ap.add_argument("--tkdiff", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--json", type=Path)
    a = ap.parse_args()
    from parity_lib import lo_recalc, set_inputs  # noqa: E402

    new, old = openpyxl.load_workbook(a.new, data_only=True), openpyxl.load_workbook(a.old, data_only=True)
    newf = openpyxl.load_workbook(a.new)
    diff = json.loads(a.tkdiff.read_text(encoding="utf-8"))
    cfg, loc = load_scenarios(), Locator(a.new)
    tmp = Path(tempfile.mkdtemp(prefix="rep_v014_"))
    books = {"v014": new, "v013_file": old}
    for sc in ("v013", "v14_only_luna", "v14_only_rest"):
        src = tmp / f"in_{sc}" / a.new.name
        src.parent.mkdir(parents=True)
        set_inputs(a.new, src, resolve_overrides(sc, loc, cfg))
        books[sc] = openpyxl.load_workbook(lo_recalc(src, tmp / f"out_{sc}"), data_only=True)
    out = openpyxl.Workbook()
    B = Font(bold=True)
    summary = {}

    # ① TK 名稱
    ws = out.active
    ws.title = "①TK名稱新舊值"
    ws.append(["名稱", "單位", "標籤", "欄數", "舊值（v19.0；基準欄）", "新值（v19.1；基準欄）", "最大變動率", "公式引用處（頁：格數）", "本模型用法", "語意（v19.1 變動原因）"])
    refs = {}
    for sh in newf.worksheets:
        if sh.title in ("TK_Link", "_Defaults"):
            continue
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for m in set(re.findall(r"\bTK_[A-Za-z0-9_]+", c.value)):
                        refs.setdefault(m, {}).setdefault(sh.title, 0)
                        refs[m][sh.title] += 1
    kinds = {"進命題計算": 0, "對照列（不回饋）": 0, "診斷（EtaDiag）": 0, "僅快照（無公式引用）": 0}
    for ch in diff["changed"]:
        nm = ch["name"]
        r = refs.get("TK_" + nm, {})
        use = "；".join(f"{k}：{v}" for k, v in r.items()) or "—"
        if nm == "IF_TokGW_Luna":
            kind = "進命題計算（Compute：Luna 層每 GW 產能）"
            kinds["進命題計算"] += 1
        elif r and set(r) <= {"EtaDiag"}:
            kind = "診斷（EtaDiag；不回饋）"
            kinds["診斷（EtaDiag）"] += 1
        elif r:
            kind = "對照列（" + "、".join(r) + "；不回饋）"
            kinds["對照列（不回饋）"] += 1
        else:
            kind = "僅快照（無公式引用）"
            kinds["僅快照（無公式引用）"] += 1
        cols = ch["cols"]
        base = next((c for c in cols if "基準" in c["hdr"] and "VR200" in c["hdr"]), None) or next((c for c in cols if "基準" in c["hdr"]), cols[0])
        rel = max((abs(c["rel"]) for c in cols if c["rel"] is not None), default=None)
        why = "上游輸入改變（Luna CSA top-k 1024 → 512 經 Luna 每 GW 產能往下傳）；定義未改（TK v19.1 為次版）"
        ws.append([nm, ch["unit"], ch["label"], ch["n"], base["old"], base["new"], rel, use, kind, why])
    ws.append([])
    ws.append(["彙總", json.dumps(kinds, ensure_ascii=False)])
    summary["kinds"] = kinds

    # ② 差異拆解
    w2 = out.create_sheet("②差異拆解")
    w2.append(["輸出", "版本"] + YEARS)
    cols_sc = [("v0.1.3（TK v19.0；檔內快取）", "v013_file"), ("v0.1.3（以 v0.1.4 檔回推 tk19_back）", "v013"), ("只換 Luna", "v14_only_luna"),
               ("只換其餘 38 名稱", "v14_only_rest"), ("v0.1.4（TK v19.1）", "v014")]
    table = {}
    for nmx, lab in KEY_OUT:
        vals = {}
        for sl, sc in cols_sc:
            vals[sc] = named(books[sc], nmx)
            w2.append([lab, sl] + vals[sc])
        base, nw = vals["v013_file"], vals["v014"]
        w2.append([lab, "v0.1.4 − v0.1.3"] + [(y - x) if num(x) and num(y) else None for x, y in zip(base, nw)])
        w2.append([lab, "　其中 Luna"] + [(y - x) if num(x) and num(y) else None for x, y in zip(base, vals["v14_only_luna"])])
        w2.append([lab, "　其中其餘 38 名稱"] + [(y - x) if num(x) and num(y) else None for x, y in zip(base, vals["v14_only_rest"])])
        w2.append([])
        table[nmx] = {sc: v for sc, v in vals.items()}
    w2.append(["標量輸出", "v0.1.3（檔內）", "v0.1.4", "差"])
    scal = {}
    for nmx, lab in SCALAR_OUT + EXP_OUT:
        try:
            if nmx.startswith("COST_PropGap_VR_Exp"):
                o, n_ = named(books["v013_file"], nmx)[4:6], named(books["v014"], nmx)[4:6]      # 2029、2030
                w2.append([lab + "（2029／2030）", str([round(x, 4) for x in o]), str([round(x, 4) for x in n_]), str([round(y - x, 4) for x, y in zip(o, n_)])])
                scal[nmx] = (o, n_)
                continue
            o, n_ = named(books["v013_file"], nmx)[0], named(books["v014"], nmx)[0]
        except Exception as e:                                          # noqa: BLE001
            w2.append([lab, f"讀取失敗：{e}"])
            continue
        w2.append([lab, o, n_, (n_ - o) if num(o) and num(n_) else ("相同" if o == n_ else "不同")])
        scal[nmx] = (o, n_)
    summary["scalars"] = {k: [list(x) if isinstance(x, list) else x for x in v] for k, v in scal.items()}
    summary["table"] = {k: {sc: v for sc, v in d.items()} for k, d in table.items()}

    # ③ η 診斷補充前後
    w3 = out.create_sheet("③η診斷補充前後")
    w3.append(["項目", "v0.1.1（TK v5.26）", "v0.1.2（v18.0）", "v0.1.3（v19.0）", "v0.1.4（v19.1）", "v0.1.4 ÷ v0.1.3", "說明"])
    sc_v = {}
    for sc in ("v011", "v012"):
        src = tmp / f"in_{sc}" / a.new.name
        src.parent.mkdir(parents=True)
        set_inputs(a.new, src, resolve_overrides(sc, loc, cfg))
        sc_v[sc] = openpyxl.load_workbook(lo_recalc(src, tmp / f"out_{sc}"), data_only=True)
    eta_tbl = {}
    for nmx, lab in ETA_ROWS:
        v11, v12, v13, v14 = (named(sc_v["v011"], nmx)[0], named(sc_v["v012"], nmx)[0], named(books["v013"], nmx)[0], named(new, nmx)[0])
        w3.append([lab, v11, v12, v13, v14, v14 / v13, "f_U'、殘差' 為 v0.1.4 新列，前幾版以同一公式在對應 TK 值下重算" if nmx in ("ETA_fU_Prod", "ETA_Resid2") else ""])
        eta_tbl[nmx] = [v11, v12, v13, v14]
    cap = lambda wb: wb["EtaDiag"]["D" + str(next(r for r in range(6, 80) if str(wb["EtaDiag"].cell(r, 2).value).startswith("本模型付費層級組合下 2025 機隊每 GW 付費產能")))].value  # noqa: E731
    capv = [cap(sc_v["v011"]), cap(sc_v["v012"]), cap(books["v013"]), cap(new)]
    w3.append(["機隊每 GW 付費產能（M tok/GW/年）"] + capv + [capv[3] / capv[2], "TK 絕對產能水準（不另設因子）"])
    eta_tbl["cap"] = capv
    w3.append([])
    ed = new["EtaDiag"]
    w3.append(["EtaDiag 第七節（v0.1.4 新增）", "單位", "基準", "低", "高", "標記", "說明"])
    for r in range(6, ed.max_row + 1):
        if ed.cell(r, 1).value and re.fullmatch(r"E(3[89]|4\d)", str(ed.cell(r, 1).value)):
            w3.append([f"{ed.cell(r, 1).value} {ed.cell(r, 2).value}", ed.cell(r, 3).value, ed.cell(r, 4).value, ed.cell(r, 5).value, ed.cell(r, 6).value, ed.cell(r, 7).value, ed.cell(r, 8).value])
    summary["eta"] = eta_tbl

    # ④ 逐格比對
    w4 = out.create_sheet("④逐格比對")
    w4.append(["比對", "頁", "格數", "數值／公式差異格數", "說明文字改寫格數", "差異的列（列號：列名，最多 12 列）"])

    def compare(x, label):
        res = []
        for sh in ("Inputs", "Demand", "Revenue", "Compute", "Cost", "Funding", "Reverse"):
            n = d = notes = 0
            rows = {}
            for row in old[sh].iter_rows():
                for c in row:
                    v = x[sh][c.coordinate].value
                    n += 1
                    if same(c.value, v):
                        continue
                    if isinstance(v, str) and isinstance(c.value, str):
                        notes += 1
                    else:
                        d += 1
                        rows.setdefault(c.row, old[sh].cell(c.row, 2).value)
            res.append([label, sh, n, d, notes, "；".join(f"{r}:{str(t)[:24]}" for r, t in list(rows.items())[:12])])
        return res
    cmp_rows = compare(new, "v0.1.4 基準 vs v0.1.3") + compare(books["v013"], "v0.1.4 回推 tk19_back vs v0.1.3") + compare(books["v14_only_rest"], "只換其餘 38 名稱 vs v0.1.3")
    for r in cmp_rows:
        w4.append(r)
    summary["compare"] = cmp_rows
    # 新增格（v0.1.4 比 v0.1.3 多出的列）：EtaDiag、Checks、TK_Link
    w4.append([])
    w4.append(["範圍外變動檢查", "EtaDiag 既有列 E01–E37 數值差異格數（v0.1.4 vs v0.1.3 檔）"])
    eo, en = old["EtaDiag"], new["EtaDiag"]
    dd = [(r, c) for r in range(6, 49) for c in (4, 5, 6) if not same(eo.cell(r, c).value, en.cell(r, c).value)]
    w4.append(["", len(dd), "；".join(f"{eo.cell(r, 1).value}" for r, _ in dd[:30])])
    summary["eta_old_rows_changed"] = sorted({eo.cell(r, 1).value for r, _ in dd})

    # ⑤ 預設
    w5 = out.create_sheet("⑤已套用的預設")
    w5.append(["問題", "預設", "替代", "對結果影響"])
    for r in DEFAULTS:
        w5.append(list(r))
    for w in out.worksheets:
        for c in w[1]:
            c.font = B
        w.column_dimensions["A"].width = 46
    out.save(a.out)
    if a.json:
        a.json.write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print("saved", a.out)


if __name__ == "__main__":
    main()
