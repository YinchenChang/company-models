#!/usr/bin/env python3
"""v0.1.7 對照 Excel 與報告數字（CLAUDE.md 第 5 節、第 8 節；工作單 20261031 v0.1.7）。

①新增或變動的列（Inputs INP_197–214、Compute 第十六節、Cost 第二十五節、Funding 第十二節、Checks C182–C192、改式的既有列）
②關鍵輸出 FY2025–2030：v0.1.6 vs v0.1.7（承諾進基準）vs 不進基準
③W1–W6 瀑布（從 v0.1.6 逐項加入）＋「只動該項」（v0.1.6 只加該項；v0.1.7 只關該項）
④已套用的預設（含替代讀法的量化）
⑤S-1 替換對照（CLAUDE.md 第 8 節第 4 步）
⑥TK 名稱變動表（v21.2 → v21.7，只列實際讀取者）
情境一律以 engine（pycel）重算現行活頁簿（覆寫值取自 tools/html_scenarios.yaml 的群組定義，與 HTML 同一套）；本程式不做模型計算。
用法：python3 tools/make_report_v017.py --out docs/reports/20261031_v0.1.7_對照.xlsx [--json build_out/report_v017.json]
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "tools"))
from build_html import Locator, load_scenarios, resolve_overrides  # noqa: E402
from engine import Engine, current_model_path, resolve_key  # noqa: E402

YEARS = [2025, 2026, 2027, 2028, 2029, 2030]
KEY = [("REV_GrossCapped", "營收總額（$B）"), ("REV_NetCapped", "營收淨額（$B）"), ("COST_PropRev_VR", "每 VR 等值 GW 營收淨額（$B/GW/年）"),
       ("COST_PropFull_VR", "每 VR 等值 GW 全成本"), ("COST_PropGap_VR", "每 VR 等值 GW 差額（$B/GW/年）"), ("COST_Coverage", "覆蓋率"),
       ("CMP_SupplyGW", "供給 GW"), ("CMP_Supply_VReq", "供給 VR 等值 GW"), ("COST_Compute", "算力成本（$B）"), ("COST_ContractPay", "合約實付（$B）"),
       ("COST_OwnedCapex", "自建資本支出（$B）"), ("COST_NonCompExSBC", "非算力成本（不含股權報酬，$B）"), ("FND_FCF", "營運自由現金流（$B）"),
       ("FND_CashNoExt", "只靠已到位融資的年底現金（$B）"), ("FND_MinCash", "最低現金（$B）"), ("FND_ExtNeedCum", "累計外部資金需求（$B）")]
SCAL = [("FND_CashNoExtTrough", "現金谷底（$B）"), ("FND_FirstGapYear", "首次缺口年"), ("CMP_Eta2025", "η（2025）"), ("FND_L1_ExtTotal", "第 1 層累計外部資金需求（$B）"),
        ("COST_S1Unalloc", "S-1 承諾未歸屬餘額（$B）"), ("COST_S1fNonComp2026", "2026 非算力成本（S-1 校準，$B）")]
STEPS = [("v0.1.6", "v016"), ("＋W1 承諾進基準", "wf_w1"), ("＋W2 自建改 GW 法", "wf_w2"), ("＋W3 2025 校準", "wf_w3"), ("＋W4 非算力校準", "wf_w4"),
         ("＋W5 快照同步（＝v0.1.7；W6 不改數值）", "base")]
ITEMS = [("W1 承諾進基準", "c_off"), ("W2 自建改 GW 法", "own_old"), ("W3 2025 校準", "cal25_off"), ("W4 非算力校準", "d22_off"), ("W5 TK v21.7", "tk16_back")]
ALTS = [("未歸屬餘額付款期 快 2027–2031", "s1u_fast"), ("未歸屬餘額付款期 慢 2028–2036", "s1u_slow"), ("SpaceXAI 以 S-1 上限 120.5 計", "spx_cap"),
        ("Google 額外 TPU 計入（中點）", "gaddl_on"), ("Broadcom 起始年 2026（S-1 期間起點）", "bcm_2026")]
DEFAULTS = [
    ("未歸屬餘額付款期", "2027–2033 平均（Assumed；Inputs s1u_start／s1u_end）", "快 2027–2031／慢 2028–2036（Sens：兩列同取低欄／高欄）", "s1u_fast", "s1u_slow"),
    ("上限型合約（SpaceXAI ≤120.5、Nscale ≤44.6）", "沿用既有規則：SpaceX＝月費 × 月數（45；上限只作區間，差額落入未歸屬餘額）；Nscale＝U2 上限全數計入（44.6）", "SpaceXAI 以上限 120.5 計（Inputs s1f_spacex_cap_switch＝1）", "spx_cap", None),
    ("Google 額外 TPU 125–150", "不進基準（Inputs s1f_gaddl_switch＝0）", "計入中點 137.5，2028–2031 平均", "gaddl_on", None),
    ("自建 GW 無證據", "0（Wattonomics v0.1.7 各年 0；未找到≠不存在，列為缺口）", "舊法 Fluidstack $50B 三年均攤（own_capex_method＝0）", "only_w2_off", None),
    ("仍為二手轉載的 S-1 列", "維持 Interested-party；不改值", "以瀏覽器核對後改 Verified", None, None),
    ("非算力成本推得值與現值差 >30%", "照推得值入基準（+82%），報告首節說明並列舊值", "維持 v0.1.6 算法（s1f_d22_switch＝0）", "only_w4_off", None),
    ("未歸屬餘額扣除範圍", "S-1 三段合計 − 模型逐家（不含 AMD：S-1 寫明 AMD「不在承諾表內」）", "連 AMD 一起扣（未歸屬餘額少 73.4）", None, None),
    ("Broadcom 付款與容量期間", "S-1 期間 2026–2031 的 2026 視為簽約年；首個爬坡年沿用 8-K 的 2027（同 Microsoft 2026-11 → 2027 的整數化）；付款改 161.2、容量維持 3.5 GW", "起始年 2026（Inputs broadcom_start 低欄）", "bcm_2026", None),
    ("AMD 合約型別", "維持 GW 型（S-1 仍只給下限 >20、無期間；既有規則）", "改金額型 20", None, None),
    ("未歸屬餘額的晶片族", "通用合約價 12 × 持有比 1＝NVIDIA VR200 族（VR 等值係數 1）", "TPU／Trainium（持有比較低 → GW 較多）", None, None),
    ("W4 下半年非算力費用", "Q2 持平（同 U1 的 Q4＝Q3 慣例；S-1 無 Q3 費用）", "上半年 × 2（16.07，較低）", None, None),
    ("W4 扣除項", "扣模型 2026 雲端平台抽成（模型自營收扣除，避免重複）與模型股權報酬（S-1 上半年股權報酬未取得）", "不扣股權報酬（非算力更高）", None, None),
    ("W2 廠房提前入帳", "lag 18 個月（同 Wattonomics lag_bldg；權重 0／0.5／0.5）；2031 年後新增不計", "12 或 24 個月", None, None),
    ("W3 2025 非算力費用", "SRC_ANT_277（草擬版 12.65 − 7.33＝5.32）不改；S-1 同法 5.319，差 0.001", "改用 S-1 推得值", None, None),
    ("W5 OpenAI 並排", "OAI_Link 改 OpenAI v0.6.4（其 TK 亦 v21.7）", "維持 v0.6", None, None),
    ("W6 Amazon", "模型 2026 已到位融資中的 Amazon＝Series G 5＋Series H 內含 5＝10.0，與 10-Q 一致 → 不改；18.0 只在回流對照累計欄", "—", None, None),
]
# CLAUDE.md 第 8 節第 4 步：本版改指正式 S-1 的引用（舊列 → 新列）
REPL = [("SRC_ANT_001", "SRC_ANT_412", "W3：Revenue 2025 校準（REV_Target 2025、REV_Gap2025、通路占比、其他營收）"),
        ("SRC_ANT_252", "SRC_ANT_424", "W3：Cost act25、Compute e_spend（η 校準）"),
        ("SRC_ANT_247", "SRC_ANT_450", "W1：Google 合約總額（金額型；付款與 GW）"),
        ("SRC_ANT_248", "SRC_ANT_451", "W1：Amazon 合約總額"),
        ("SRC_ANT_211", "SRC_ANT_452", "W1：Microsoft 合約總額"),
        ("SRC_ANT_250", "SRC_ANT_454", "W1：Broadcom 付款總額（雙揭露型；容量不變）"),
        ("SRC_ANT_400", "SRC_ANT_455", "W1：Nscale 錨定總額（U2）"),
        ("SRC_ANT_400", "SRC_ANT_456", "W1：Lambda 錨定總額（U2）")]


def named_vals(eng, n):
    v = eng.get_name(n)
    return v if isinstance(v, list) else [v]


def run(eng, loc, cfg, ov_scn=None, ov=None):
    ov = dict(ov or {})
    if ov_scn:
        ov.update(resolve_overrides(ov_scn, loc, cfg) if ov_scn in cfg["scenarios"] else {})
    orig = {k: eng.get(*resolve_key(loc.names, k)) for k in ov}
    try:
        for k, v in ov.items():
            eng.set_key(k, v)
        out = {n: named_vals(eng, n) for n, _ in KEY}
        out.update({n: eng.get_name(n) for n, _ in SCAL})
        out["CHK_Errors"] = eng.get_name("CHK_Errors")
        return out
    finally:
        for k, v in orig.items():
            eng.set_key(k, v)


def group_ov(cfg, loc, groups):
    tmp = {"scenarios": {"_t": groups}, "groups": cfg["groups"]}
    return resolve_overrides("_t", loc, tmp)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--json", type=Path, default=REPO / "build_out" / "report_v017.json")
    a = ap.parse_args()
    model = current_model_path()
    loc, cfg, eng = Locator(model), load_scenarios(), Engine(model)
    R = {}
    for _l, sc in STEPS + [("", "commit_off")] + [("", s) for _l2, s in ALTS] + [("", "only_w2_off"), ("", "only_w4_off")]:
        R[sc] = run(eng, loc, cfg, sc if sc != "base" else None)
    off_all = ["c_off", "own_old", "cal25_off", "d22_off", "tk16_back"]
    R["only_on"] = {}
    for lab, g in ITEMS:
        rest = [x for x in off_all if x != g]
        R["only_on"][lab] = run(eng, loc, cfg, ov=group_ov(cfg, loc, rest))
    R["only_off"] = {lab: run(eng, loc, cfg, ov=group_ov(cfg, loc, [g])) for lab, g in ITEMS}

    wb = openpyxl.Workbook()
    B = Font(bold=True)

    def sheet(title, header):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = B
        return ws

    # ① 新增與變動列
    ws = sheet("①新增與變動列", ["頁", "列 ID", "列名", "D", "E", "F", "G", "H", "I", "來源／標記／說明"])
    mv = openpyxl.load_workbook(model, data_only=True)
    sections = {"Compute": "十六、", "Cost": "二十五、", "Funding": "十二、"}
    for sh, pre in sections.items():
        w = mv[sh]
        start = next(r for r in range(1, w.max_row + 1) if str(w.cell(r, 1).value or "").startswith(pre))
        for r in range(start, w.max_row + 1):
            if w.cell(r, 2).value is None and w.cell(r, 1).value is None:
                continue
            ws.append([sh, w.cell(r, 1).value, w.cell(r, 2).value] + [w.cell(r, c).value for c in range(4, 10)] + [w.cell(r, 11).value])
    wi = mv["Inputs"]
    for r in range(5, wi.max_row + 1):
        iid = wi.cell(r, 1).value
        if isinstance(iid, str) and iid.startswith("INP_") and int(iid[4:]) >= 197:
            ws.append(["Inputs", iid, wi.cell(r, 2).value, wi.cell(r, 5).value, wi.cell(r, 6).value, wi.cell(r, 7).value, None, None, None,
                       f"{wi.cell(r, 8).value}｜{wi.cell(r, 9).value}"])
    ws.append(["Inputs", "INP_131", "broadcom_start 低欄 2027 → 2026（正式 S-1 期間起點；替代讀法）", 2027, 2026, 2027, None, None, None, "Decision"])
    wc = mv["Checks"]
    for r in range(5, wc.max_row + 1):
        cid = wc.cell(r, 1).value
        if isinstance(cid, str) and cid.startswith("C") and cid[1:].isdigit() and int(cid[1:]) >= 182:
            ws.append(["Checks", cid, wc.cell(r, 2).value, wc.cell(r, 3).value, wc.cell(r, 4).value, wc.cell(r, 5).value, None, None, None, wc.cell(r, 6).value])
    for sh, rid, lab in (("Compute", "—", "own_capex：加 own_capex_method 開關（GW 法／舊法）"), ("Compute", "—", "sup_own：同上"),
                         ("Compute", "—", "sup_con：加未歸屬餘額、Google 額外 TPU"), ("Compute", "—", "gw_vr200／gw_tpu7：加未歸屬餘額（VR200）、Google 額外 TPU（TPU v7）"),
                         ("Compute", "—", "e_spend：加 s1f_cal25_switch"), ("Cost", "—", "act25：加 s1f_cal25_switch"), ("Cost", "—", "pay：加未歸屬餘額、Google 額外 TPU"),
                         ("Cost", "—", "pay_spacex、tot_spacex：× SpaceXAI 上限讀法係數（基準 1）"), ("Cost", "—", "tot_google／broadcom／aws／azure：T＝開關 × 正式 S-1 ＋（1 − 開關）× 舊列"),
                         ("Cost", "—", "u2_T_n／u2_T_l：正式 S-1 Nscale 44.6／Lambda 41.3（開關）"), ("Cost", "—", "pp（每人年成本）2026 欄：s1f_d22_switch"),
                         ("Revenue", "—", "2025 校準營收 5 處：s1f_cal25_switch"), ("Reverse", "—", "2025 目標：s1f_cal25_switch"),
                         ("Funding", "—", "FND_CommitAfter2030：加未歸屬餘額與 Google 額外 TPU"), ("Checks", "C109", "U2 總額核對：承諾開關 1 時對正式 S-1 44.6＋41.3"),
                         ("TK_Link", "—", "Tokenomics v21.7（master 90f7deb）；新增 IF_CapexIT、IF_CapexFacility"), ("OAI_Link", "—", "OpenAI v0.6.4")):
        ws.append([sh, rid, "改式／更新：" + lab])

    # ② 關鍵輸出
    ws = sheet("②關鍵輸出", ["輸出", "讀法"] + [str(y) for y in YEARS])
    for n, lab in KEY:
        for rl, sc in (("v0.1.6", "v016"), ("v0.1.7（承諾進基準）", "base"), ("不進基準（W1 關閉）", "commit_off")):
            ws.append([lab, rl] + R[sc][n])
    for n, lab in SCAL:
        for rl, sc in (("v0.1.6", "v016"), ("v0.1.7（承諾進基準）", "base"), ("不進基準（W1 關閉）", "commit_off")):
            ws.append([lab, rl, R[sc][n]])

    # ③ 瀑布
    ws = sheet("③瀑布", ["輸出"] + [l for l, _s in STEPS])
    rows3 = [(f"每 VR 等值 GW 差額 {y}", "COST_PropGap_VR", YEARS.index(y)) for y in YEARS[1:]] + \
            [("覆蓋率 2030", "COST_Coverage", 5), ("累計外部資金需求 2030", "FND_ExtNeedCum", 5), ("現金谷底", "FND_CashNoExtTrough", None),
             ("η（2025）", "CMP_Eta2025", None), ("算力成本 2028", "COST_Compute", 3), ("非算力成本 2026", "COST_NonCompExSBC", 1)]
    for lab, n, i in rows3:
        ws.append([lab] + [(R[sc][n][i] if i is not None else R[sc][n]) for _l, sc in STEPS])
    ws.append([])
    ws.append(["只動該項：v0.1.6 只加該項（其餘維持 v0.1.6）"] + [lab for lab, _g in ITEMS])
    for lab, n, i in rows3:
        ws.append([lab] + [(R["only_on"][it][n][i] if i is not None else R["only_on"][it][n]) for it, _g in ITEMS])
    ws.append([])
    ws.append(["只動該項：v0.1.7 只關該項（其餘維持 v0.1.7）"] + [lab for lab, _g in ITEMS])
    for lab, n, i in rows3:
        ws.append([lab] + [(R["only_off"][it][n][i] if i is not None else R["only_off"][it][n]) for it, _g in ITEMS])

    # ④ 已套用的預設
    ws = sheet("④已套用的預設", ["問題", "採用的預設", "替代", "替代：2027 差額", "替代：2030 差額", "替代：累計外部資金", "替代：現金谷底", "第二替代：累計外部資金", "第二替代：現金谷底"])
    for q, d, alt, s1, s2 in DEFAULTS:
        row = [q, d, alt]
        if s1:
            o = R[s1]
            row += [o["COST_PropGap_VR"][2], o["COST_PropGap_VR"][5], o["FND_ExtNeedCum"][5], o["FND_CashNoExtTrough"]]
        else:
            row += [None] * 4
        if s2:
            row += [R[s2]["FND_ExtNeedCum"][5], R[s2]["FND_CashNoExtTrough"]]
        ws.append(row)

    # ⑤ S-1 替換對照
    ws = sheet("⑤S-1替換對照", ["舊列", "舊值", "新列", "新值", "差異比例", "受影響（本版改指新列）", "命題輸出前（v0.1.6）→後（v0.1.7）", "結論是否翻轉"])
    sv = {}
    wsrc = mv["SRC_ANT"]
    for r in range(5, wsrc.max_row + 1):
        if str(wsrc.cell(r, 1).value or "").startswith("SRC_ANT_"):
            sv[wsrc.cell(r, 1).value] = wsrc.cell(r, 5).value if wsrc.cell(r, 5).value is not None else (wsrc.cell(r, 6).value or wsrc.cell(r, 7).value)
    g16, g17 = R["v016"], R["base"]
    summ = (f"2030 差額 {g16['COST_PropGap_VR'][5]:+.2f}→{g17['COST_PropGap_VR'][5]:+.2f}；2030 覆蓋率 {g16['COST_Coverage'][5]:.3f}→{g17['COST_Coverage'][5]:.3f}；"
            f"累計外部資金 {g16['FND_ExtNeedCum'][5]:.1f}→{g17['FND_ExtNeedCum'][5]:.1f}；現金谷底 {g16['FND_CashNoExtTrough']:.1f}→{g17['FND_CashNoExtTrough']:.1f}")
    for old, new, aff in REPL:
        o, nv = sv.get(old), sv.get(new)
        diff = (nv - o) / o if isinstance(o, (int, float)) and isinstance(nv, (int, float)) and o else None
        ws.append([old, o, new, nv, diff, aff, summ, "部分：命題 1 轉正年 2029→2030（2029 由正轉負）；命題 2 仍為 0、覆蓋率仍跨過 1（2030）"])

    # ⑥ TK 名稱變動
    ws = sheet("⑥TK名稱變動", ["名稱", "格數", "v21.2（v0.1.6）", "v21.7（v0.1.7）", "是否進計算"])
    old = openpyxl.load_workbook(REPO / "model" / "archive" / "20261031_Anthropic_v0.1.6.xlsx", data_only=True)
    into = {"TK_NNV_TPUv7_Hold", "TK_NNV_Trn3_Hold", "TK_NNV_MI455X_Hold", "TK_NNV_MI455X_Out", "TK_IF_HoldEcon", "TK_IF_CapexTotal"}
    for n in cfg["groups"]["tk16_back"]["set"]:
        def vals(wbk):
            sh, rg = wbk.defined_names[n].attr_text.replace("$", "").split("!")
            c = wbk[sh][rg]
            return [x.value for r_ in c for x in r_] if isinstance(c, tuple) else [c.value]
        vo, vn = vals(old), vals(mv)
        ch = [(k, p, q) for k, (p, q) in enumerate(zip(vo, vn)) if p != q]
        note = {"TK_NNV_TPUv7_Hold": "是（TPU 合約價、既有雲端加權合約價 → η、GW）", "TK_NNV_Trn3_Hold": "是（Trainium 合約價 → η、GW）",
                "TK_NNV_MI455X_Hold": "是（AMD 合約價 → 付款）", "TK_NNV_MI455X_Out": "是（AMD VR 等值係數）",
                "TK_IF_HoldEcon": "參考（供應商持有成本，只低成本欄變）", "TK_IF_CapexTotal": "否（只低成本欄變；本模型用基準欄）"}.get(n, "否（只對照或快照；高低區間格不進基準）")
        ws.append([n, len(ch), "; ".join(f"[{k + 1}] {p}" for k, p, _q in ch[:4]), "; ".join(f"[{k + 1}] {q}" for k, _p, q in ch[:4]), note])

    del wb["Sheet"]
    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    a.json.parent.mkdir(parents=True, exist_ok=True)
    json.dump(R, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
    print("saved", a.out, a.json)


if __name__ == "__main__":
    main()
