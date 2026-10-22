#!/usr/bin/env python3
"""v0.1.3 B 段對照 Excel（CLAUDE.md 第 5 節）：①新增與變動列 ②前後對照與逐項拆解 ③η 拆解表 ④外部驗收報告逐項處理 ⑤已套用的預設。

用法：python3 tools/make_report_v013b.py --new model/20261022_Anthropic_v0.1.3.xlsx --v012 model/archive/20261021_Anthropic_v0.1.2.xlsx \
        --v011 model/archive/20261021_Anthropic_v0.1.1.xlsx --tk-dir /tmp/tk --out docs/reports/20261022_v0.1.3_對照.xlsx
數值一律讀 LibreOffice 重算後的快取值；本程式不做模型計算（版本變動小表的因子由各版快取格與 Tokenomics 檔計算，屬報告層的診斷）。
另以 u2_method＝0 重算一份，與 v0.1.2 逐格比對（證明 U2 以外沒有任何變動）。
"""
import argparse
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "tests" / "parity"))

BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="FFEDEDED")
YEARS = ["2025", "2026", "2027", "2028", "2029", "2030"]
KEY_OUT = [
    ("COST_PropGap_VR", "每 VR 等值 GW 差額（含股權報酬，$B/GW/年）"), ("COST_PropRev_VR", "每 VR 等值 GW 營收淨額（$B/GW/年）"),
    ("COST_PropFull_VR", "每 VR 等值 GW 全成本（$B/GW/年）"), ("COST_Coverage", "覆蓋率（營收淨額 ÷ 全成本）"),
    ("CMP_Supply_VReq", "供給 VR 等值 GW"), ("CMP_SupplyGW", "供給 GW"), ("COST_Compute", "算力成本（現金，$B）"), ("REV_NetCapped", "營收淨額（截頂後，$B）"),
    ("FND_FCF", "營運自由現金流（$B）"), ("FND_CashEnd", "年底現金（基準，$B）"), ("FND_ExtNeedCum", "累計外部資金需求（基準，$B）"), ("CMP_Eta", "η"),
]
DEFAULTS = [
    ("f_P 計價口徑", "每 GW 合約實付 ÷ TK 持有成本（經濟口徑，加權）；區間＝TK 低／高成本情境", "反向（持有成本 ÷ 合約實付）", "只是方向不同；TK 方向的倒數並列於 EtaDiag 第六節"),
    ("f_U 利用率", "TK IF_UtilActual（0.41；0.30–0.65）÷ IF_Util（0.60），Analogy；不計 TK 的生產折減 CTL_ProdDerate", "計入生產折減 0.85（f_U＝0.581）", "計入則 f_U −15%，殘差 +18%（0.446 → 約 0.525）；本模型沒有 Anthropic 自己的利用率資料"),
    ("f_W 工作負載（層級）組合", "機隊加權：逐族把本模型組合後產能換成 TK 參考組合（B4_MixPaid）產能，TPU v6e／Trainium2 隨 Hopper、TPU v7／Trainium3／MI455X 隨 VR200", "只用 VR200 近似（f_W＝1.390）", "機隊加權 1.399 對 1.390（+0.7%）；2025 機隊是 Hopper／TPU v6e／Trainium2，故機隊加權為正式口徑"),
    ("輸入輸出組合", "標「無法分離」並入殘差（每任務 token 已直接取自 TK 任務模板，含新鮮、快取、decode）", "把 TK L1_GapPrompt（1.305）× L1_GapISL（1.0）當 Analogy 因子", "若採用，殘差 +30%（0.446 → 約 0.58）；但那是 OpenAI 口徑的每則提示 token，不是 Anthropic 的"),
    ("Bridge 低／高", "2030 敏感度包絡（Compute 第十四節 A 合約價 8／20、B η 沿用／回升、C 產出比低／高）取最小／最大；2025 每任務 token 用 Inputs 倍數低／高", "全部用 Inputs 低／高欄重算", "營收側與用戶側 Inputs 區間未另做重算列，仍以 tools/sensitivity.py 與 HTML ⑥ 為準"),
    ("Bridge 與 TK 對照年度", "2026（TK 的 P × Q 以 2026 Q2 年化 m 為主口徑）", "2025", "TK 無 2025 的 P × Q 對應列"),
    ("EtaDiag 另立一頁", "獨立唯讀頁（不放在 Compute 尾端）", "放在 Compute 尾端", "Compute 是計算頁，診斷放在計算頁會模糊「不回饋」的界線；Checks e1_iso＋test_bridge_not_fed_back 保證隔離"),
    ("版本變動小表的 f_U", "TK v5.26 沒有 IF_UtilActual，三個版本一律用 0.41／0.60 比較", "v5.26 用另一假設值", "f_U 在三版相同，不影響「版本變動被哪一項吸收」的判讀"),
    ("TK_Link 新增 23 名的位置", "放在 NonNV 讀表之後（頁尾）", "接在 v0.1.2 的 14 名之後", "避免移動既有列"),
    ("HTML 三小節的位置", "併入「①之三」並改名「v0.1.3 更新了什麼」（資金三層、持續擴張在 A 段已在該節）", "另開 ①之四", "避免首頁出現兩個『更新了什麼』"),
    ("OpenAI 並排快照", "維持 v0.6（取 archive/20261008_OpenAI_v0.6.xlsx；SHA-256 34fc34a7…）；main 現行 v0.6.2", "改讀 v0.6.2", "只影響並排列；不在本版範圍"),
]
ACCEPT = [
    ("Priority 4：資金命題拆三層", "採納", "A", "Funding 第十節 FND_L1/2/3_*；HTML ①之三", "第 1 層（營運自給）2026–2030 累計仍需外部提供 41.9；第 2 層（已到位股權＝基準）0；第 3 層 0"),
    ("Priority 3：持續擴張情境（2029–2030 追加算力）", "採納（只作情境，基準仍只計已簽合約）", "A", "Cost 第二十三節、Funding 第十一節；Inputs INP_189–194", "無需求／部分變現／完全變現三檔；基準不受影響"),
    ("結論措辭改條件式", "採納", "A", "tools/build_html.py ① 結論句、⑥「要成立必須」", "「按目前已到位融資與基準支出路徑，不出現最低現金缺口」；不寫「不需要再融資」"),
    ("Priority 2：拆分 η", "部分採納（以可得資料拆成計價口徑、利用率、工作負載組合與殘差；並列 TK L1_Gap*；不改 η 的值）", "B", "EtaDiag 頁；HTML ①之三；本報告 ③", "η 0.527 ＝ 1.237 × 0.683 × 1.399 × 0.446；輸入輸出組合無法分離；殘差吸收了 TK v5.26→v18.0 的整體變動"),
    ("Priority 1：需求 → token → GW → 成本橋接頁", "部分採納（新增唯讀 Bridge 頁；不以與 TK 逐項閉合為驗收條件）", "B", "Bridge 頁；HTML ①之三", "109 列，逐層列來源、單位、2025–2030、低／高、標記；與 TK 對照 2026 P × Q、Q、token 帳；TK token 帳 6.26 倍未閉合屬 Tokenomics 待辦"),
    ("Priority 5：公開 S-1 後的替換程序", "採納", "B", "anthropic/CLAUDE.md 第 8 節", "新增 SRC 列、舊列標已被取代並保留、模型引用改指新列、報告列舊值／新值／差異比例／受影響輸出／結論是否翻轉"),
    ("TK 快照標記 v19.0（chat 端加）", "採納", "A", "TK_Link；check_tk_snapshot", "103 個值與 v18.0 相同，數值差異 0"),
    ("U2 新法（v0.1.1 Andy 同意的預設）", "採納", "A", "Inputs u2_method；Compute／Cost 新節", "供給取揭露容量，付款路徑不變；2030 差額 11.139 → 11.602"),
    ("情境機率 50／30／20", "不採納", "—", "—", "無依據的 Assumed；模型不輸出機率"),
    ("以 TK 閉合作為『通過』門檻", "不採納", "—", "—", "屬 Tokenomics 範圍；Bridge 只並列差額"),
]


def named(wb, n):
    sh, rg = wb.defined_names[n].attr_text.split("!")
    c = wb[sh.strip("'")][rg.replace("$", "")]
    if not isinstance(c, tuple):
        c = ((c,),)
    return [x.value for r in c for x in r]


def tk_names(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    def f(n):
        sh, rg = wb.defined_names[n].attr_text.replace("$", "").split("!")
        return [c.value for row in wb[sh][rg] for c in row] if ":" in rg else [wb[sh][rg].value]
    return f


def version_factors(model, tkfile):
    """版本變動小表：由各版模型的快取格與 Tokenomics 檔計算 η 的各因子（機隊加權的 f_W；f_U 三版一律 0.41 ÷ IF_Util）。"""
    m = openpyxl.load_workbook(model, data_only=True)
    c, k = m["Compute"], m["Cost"]
    t = tk_names(tkfile)
    eta, price, hold = c["D149"].value, c["D67"].value, k["D70"].value / c["D90"].value
    u0 = t("IF_Util")[0]
    fu = 0.41 / u0
    mix = t("B4_MixPaid")                                                  # Luna、Sol、Astra
    mm = [c["D113"].value, c["D114"].value, c["D115"].value]               # 本模型付費層級組合：頂、中、低
    ref = [mix[2], mix[1], mix[0]]                                         # 同順序：頂、中、低
    ratio = {}
    for g, r0 in (("hopper", 13), ("gb200", 16), ("gb300", 19), ("vr200", 22)):
        tk3 = [c.cell(r0 + i, 4).value for i in range(3)]                  # 頂、中、低 每 GW 總產出
        ratio[g] = sum(a / b for a, b in zip(mm, tk3)) / sum(a / b for a, b in zip(ref, tk3))
    fam_gen = ["hopper", "gb200", "gb300", "vr200", "hopper", "vr200", "hopper", "vr200", "vr200"]
    s = [c.cell(102 + i, 4).value for i in range(9)]
    cap = [c.cell(120 + i, 4).value for i in range(9)]
    capmod = c["D129"].value
    capref = sum(si * ci * ratio[g] for si, ci, g in zip(s, cap, fam_gen))
    fp, fw = price / hold, capref / capmod
    prod = fp * fu * fw
    return dict(eta=eta, price=price, hold=hold, fp=fp, u0=u0, fu=fu, capref=capref, capmod=capmod, fw=fw, prod=prod, res=eta / prod,
                tokgw=c["D145"].value, spendgw=c["D148"].value, fleet_cap=capmod)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--new", required=True, type=Path)
    ap.add_argument("--v012", required=True, type=Path)
    ap.add_argument("--v011", required=True, type=Path)
    ap.add_argument("--tk-dir", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    a = ap.parse_args()
    new, o12 = openpyxl.load_workbook(a.new, data_only=True), openpyxl.load_workbook(a.v012, data_only=True)
    out = openpyxl.Workbook()

    # ① 新增與變動列
    ws = out.active
    ws.title = "①新增與變動列"
    ws.append(["頁", "編號", "項目", "單位"] + YEARS + ["低", "高", "標記", "來源／說明"])
    b = new["Bridge"]
    for r in range(7, b.max_row + 1):
        if b.cell(r, 1).value and str(b.cell(r, 1).value).startswith("B"):
            ws.append(["Bridge", b.cell(r, 1).value, b.cell(r, 2).value, b.cell(r, 3).value] + [b.cell(r, c).value for c in range(4, 10)] +
                      [b.cell(r, 12).value, b.cell(r, 13).value, b.cell(r, 10).value, f"{b.cell(r, 11).value or ''}｜{b.cell(r, 14).value or ''}"])
    ws.append([])
    ws.append(["頁", "編號", "項目", "單位", "基準", "低", "高", "", "", "", "", "", "", "標記", "說明"])
    e = new["EtaDiag"]
    for r in range(6, e.max_row + 1):
        if e.cell(r, 1).value and str(e.cell(r, 1).value).startswith("E"):
            ws.append(["EtaDiag", e.cell(r, 1).value, e.cell(r, 2).value, e.cell(r, 3).value, e.cell(r, 4).value, e.cell(r, 5).value, e.cell(r, 6).value,
                       "", "", "", "", "", "", e.cell(r, 7).value, e.cell(r, 8).value])
    ws.append([])
    ws.append(["TK_Link（新增）", "名稱", "說明", "單位", "值1", "值2", "值3"])
    ot = {r[0].value for r in o12["TK_Link"].iter_rows(min_row=10) if r[0].value}
    for r in new["TK_Link"].iter_rows(min_row=10):
        if r[0].value and r[0].value not in ot:
            ws.append(["TK_Link", r[0].value, r[2].value, r[3].value] + [r[i].value for i in (9, 10, 11)])
    ws.append([])
    ws.append(["Checks（新增）", "ID", "檢查項", "值", "期望", "狀態"])
    oc = {r[0].value for r in o12["Checks"].iter_rows(min_row=5) if r[0].value}
    for r in new["Checks"].iter_rows(min_row=5):
        if r[0].value and r[0].value not in oc and str(r[0].value).startswith("C"):
            ws.append(["Checks", r[0].value, r[1].value, r[2].value, r[3].value, r[4].value])

    # ② 前後對照與逐項拆解
    w2 = out.create_sheet("②前後對照與逐項拆解")
    w2.append(["輸出", "版本"] + YEARS)
    for n_, lab in KEY_OUT:
        ov, nv = named(o12, n_), named(new, n_)
        w2.append([lab, "v0.1.2"] + ov)
        w2.append([lab, "v0.1.3"] + nv)
        w2.append([lab, "差（新 − 舊）"] + [(y - x) if isinstance(x, (int, float)) and isinstance(y, (int, float)) else None for x, y in zip(ov, nv)])
    w2.append([])
    # 逐格比對：v0.1.3 對 v0.1.2（基準），與 u2_method＝0 的 v0.1.3 對 v0.1.2
    from parity_lib import lo_recalc, set_inputs  # noqa: E402
    tmp = Path(tempfile.mkdtemp(prefix="rep_b_"))
    src = tmp / "in" / a.new.name
    src.parent.mkdir(parents=True)
    id_u2 = None
    wf = openpyxl.load_workbook(a.new)
    for r in wf["Inputs"].iter_rows(min_row=5):
        if r[13].value == "u2_method":
            id_u2 = r[0].value
    set_inputs(a.new, src, {id_u2: 0})
    off = openpyxl.load_workbook(lo_recalc(src, tmp / "out"), data_only=True)

    def compare(x, label):
        stat = []
        for sh in ("Inputs", "Demand", "Revenue", "Compute", "Cost", "Funding", "Reverse"):
            n = d = notes = 0
            rows = {}
            for row in o12[sh].iter_rows():
                for c in row:
                    v = x[sh][c.coordinate].value
                    n += 1
                    same = (c.value == v) or (isinstance(v, (int, float)) and isinstance(c.value, (int, float)) and abs(v - c.value) <= 1e-9 * max(1, abs(v)))
                    if not same and c.column_letter == "K" and isinstance(v, str):
                        notes += 1                                     # 說明文字（K 欄）改寫，不是數值或公式
                    elif not same:
                        d += 1
                        rows.setdefault(c.row, o12[sh].cell(c.row, 2).value)
            stat.append((label, sh, n, d, notes, "；".join(f"{r}:{str(t)[:24]}" for r, t in list(rows.items())[:12])))
        return stat
    w2.append(["逐格比對（對象：v0.1.2 的 Inputs、Demand、Revenue、Compute、Cost、Funding、Reverse 全部儲存格）", "", "", "", "", ""])
    w2.append(["情境", "頁", "比對格數", "數值／公式差異格數", "說明文字（K 欄）改寫格數", "差異的列（列號：列名，最多 12 列）"])
    for rec in compare(new, "v0.1.3 基準 vs v0.1.2") + compare(off, "v0.1.3（u2_method＝0）vs v0.1.2"):
        w2.append(list(rec))
    w2.append([])
    w2.append(["變動項目", "類別", "是否影響基準輸出", "說明"])
    for rec in [
        ("TK 快照標記 v18.0 → v19.0", "TK 標記", "否", "103 個值逐一相同（check_tk_snapshot：OK 103、DIFF 0、MISSING 0）；A 段重建後 27,934 格比對數值差異 0"),
        ("U2 新法（Nscale、Lambda 供給取揭露容量）", "U2", "是（供給 GW、VR 等值 GW、命題差額）", "2030 差額 +11.139 → +11.602；成本、營收、現金流不變；u2_method＝0 時與 v0.1.2 逐格相同（上表）"),
        ("資金三層", "只作情境或顯示", "否", "新增 Funding 第十節；第 2 層＝基準（Checks d1_l2_eq）"),
        ("持續擴張情境", "只作情境或顯示", "否", "新增 Cost 第二十三節、Funding 第十一節；基準只計已簽合約"),
        ("條件式結論措辭", "只作情境或顯示", "否", "HTML 文字"),
        ("Bridge 頁", "只作情境或顯示", "否", "唯讀；builder 掃描與 test_bridge_not_fed_back 確認無計算頁引用"),
        ("EtaDiag 頁（η 拆解）", "只作情境或顯示", "否", "η 的值不變（0.52729641508883）；test_eta_decomposition_does_not_change_eta"),
        ("TK_Link 尾端 +23 名", "只作情境或顯示", "否", "IF_UtilActual（含低高）、L1_TokGW_PaidCap_VR200、B4_MixPaid、L1_Gap*（各含低高）；只被 EtaDiag 引用"),
        ("S-1 替換程序", "文件", "否", "anthropic/CLAUDE.md 第 8 節"),
    ]:
        w2.append(list(rec))

    # ③ η 拆解表
    w3 = out.create_sheet("③η拆解表")
    w3.append(["本版 η（2025）拆解（EtaDiag）", "單位", "基準", "低", "高", "標記", "說明"])
    for r in range(6, e.max_row + 1):
        if e.cell(r, 1).value and str(e.cell(r, 1).value).startswith("E"):
            w3.append([f"{e.cell(r, 1).value} {e.cell(r, 2).value}", e.cell(r, 3).value, e.cell(r, 4).value, e.cell(r, 5).value, e.cell(r, 6).value,
                       e.cell(r, 7).value, e.cell(r, 8).value])
    w3.append([])
    w3.append(["版本變動小表：TK v5.26 → v18.0 → v19.0（因子由各版模型快取格與 Tokenomics 檔計算；f_W 為機隊加權）", "v0.1.1（TK v5.26）", "v0.1.2（TK v18.0）", "v0.1.3（TK v19.0）",
               "v0.1.2 ÷ v0.1.1", "v0.1.3 ÷ v0.1.2", "說明"])
    tkd = a.tk_dir / "model"
    vs = [version_factors(a.v011, tkd / "archive" / "20261007_Tokenomics_v5.26.xlsx"),
          version_factors(a.v012, tkd / "archive" / "20261021_Tokenomics_v18.0.xlsx"),
          version_factors(a.new, tkd / "20261022_Tokenomics_v19.0.xlsx")]
    labels = [("eta", "η（2025）", "被拆解的量"), ("fp", "f_P 計價口徑", "合約實付 ÷ TK 持有成本（加權）；持有成本隨 TK 版本略降"),
              ("fu", "f_U 利用率", "三版一律 0.41 ÷ IF_Util（v5.26 無實際利用率名稱，沿用）"), ("fw", "f_W 工作負載組合（機隊加權）", "TK 參考組合產能 ÷ 本模型組合產能；兩者同比例隨 TK 產能上升，故幾乎不變"),
              ("prod", "已解釋部分 f_P × f_U × f_W", ""), ("res", "未解釋殘差", "吸收了版本變動")]
    for k_, lab, note in labels:
        w3.append([lab] + [v[k_] for v in vs] + [vs[1][k_] / vs[0][k_], vs[2][k_] / vs[1][k_], note])
    w3.append(["（對照）本模型機隊每 GW 付費產能 M tok/GW/年"] + [v["capmod"] for v in vs] + [vs[1]["capmod"] / vs[0]["capmod"], vs[2]["capmod"] / vs[1]["capmod"],
                                                                         "TK 每 GW 產能上升 → token 換算 GW 下降 → η 下降"])
    w3.append(["（對照）token 換算推論 GW（2025）"] + [v["tokgw"] for v in vs] + [vs[1]["tokgw"] / vs[0]["tokgw"], vs[2]["tokgw"] / vs[1]["tokgw"], ""])
    w3.append(["（對照）支出換算推論 GW（2025）"] + [v["spendgw"] for v in vs] + [vs[1]["spendgw"] / vs[0]["spendgw"], vs[2]["spendgw"] / vs[1]["spendgw"], "與 TK 版本無關（支出 ÷ 合約價）"])
    w3.append([])
    w3.append(["結論：TK v5.26 → v18.0 使 η 降至原來的 %.1f%%；其中 f_P 變動 ×%.3f、f_W ×%.3f、f_U ×1，殘差 ×%.3f——版本變動幾乎全被殘差吸收（因為三個因子都是比值，TK 產能整體上升不會改變它們）。v18.0 → v19.0：全部因子與 η 不變。"
               % (100 * vs[1]["eta"] / vs[0]["eta"], vs[1]["fp"] / vs[0]["fp"], vs[1]["fw"] / vs[0]["fw"], vs[1]["res"] / vs[0]["res"])])
    print("version factors:", [{k: round(v, 4) for k, v in x.items() if k in ("eta", "fp", "fu", "fw", "prod", "res")} for x in vs])

    # ④ 外部驗收報告逐項處理
    w4 = out.create_sheet("④外部驗收逐項處理")
    w4.append(["報告項", "處理", "段", "位置", "結果"])
    for rec in ACCEPT:
        w4.append(list(rec))

    # ⑤ 已套用的預設
    w5 = out.create_sheet("⑤已套用的預設")
    w5.append(["問題", "採用的預設", "替代", "對結果的影響"])
    for d in DEFAULTS:
        w5.append(list(d))

    for w in out.worksheets:
        for c in w[1]:
            c.font, c.fill = BOLD, HDR
        w.column_dimensions["A"].width = 48
        w.column_dimensions["B"].width = 40
        w.column_dimensions["C"].width = 40
        for col in "DEFGHIJKLMN":
            w.column_dimensions[col].width = 16
        for row in w.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(c.value, float):
                    c.number_format = "#,##0.000"
    a.out.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
