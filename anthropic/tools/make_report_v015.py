#!/usr/bin/env python3
"""v0.1.5 對照 Excel：①TK 10 名稱新舊值與使用位置 ②關鍵輸出 v0.1.4 vs v0.1.5（2025–2030 逐年）③逐格比較（全頁；範圍外差異）。
用法：python3 tools/make_report_v015.py --new model/20261024_Anthropic_v0.1.5.xlsx --old model/archive/20261022_Anthropic_v0.1.4.xlsx \
        --tkdiff /tmp/tkdiff.json --out docs/reports/20261024_v0.1.5_對照.xlsx [--json out.json]
數值一律讀 LibreOffice 重算後的快取值；本程式不做模型計算。
"""
import argparse
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

KEY_OUT = [("COST_PropGap_VR", "每 VR 等值 GW 差額（$B/GW/年）"), ("COST_Coverage", "覆蓋率"), ("CMP_Supply_VReq", "供給 VR 等值 GW"), ("CMP_SupplyGW", "供給 GW"),
           ("CMP_Eta", "η"), ("CMP_InfGW_Eff", "有效推論 GW"), ("CMP_RDGW", "研發 GW"), ("COST_Compute", "算力成本（$B）"), ("REV_NetCapped", "營收淨額（$B）"),
           ("FND_FCF", "營運自由現金流（$B）"), ("FND_CashEnd", "年底現金（$B）"), ("FND_ExtNeedCum", "累計外部資金需求（$B）")]
SCALARS = [("FND_L1_ExtTotal", "第 1 層累計外部資金需求"), ("FND_L2_ExtTotal", "第 2 層"), ("FND_L3_ExtTotal", "第 3 層"), ("FND_L1_Trough", "第 1 層現金谷底"),
           ("FND_CashNoExtTrough", "只靠已到位融資的現金谷底"), ("ETA_Eta", "η（2025）"), ("ETA_Resid", "殘差（不計生產折減）"), ("ETA_Resid2", "殘差（計入生產折減）")]


def named(wb, name):
    sh, rg = wb.defined_names[name].attr_text.replace("$", "").split("!")
    c = wb[sh.strip("'")][rg]
    if not isinstance(c, tuple):
        c = ((c,),)
    return [x.value for r in c for x in r]


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def same(a, b):
    return a == b or (num(a) and num(b) and abs(a - b) <= 1e-9 * max(1, abs(a), abs(b)))


def main():
    ap = argparse.ArgumentParser()
    for k in ("new", "old", "tkdiff", "out"):
        ap.add_argument("--" + k, required=True, type=Path)
    ap.add_argument("--json", type=Path)
    a = ap.parse_args()
    nv, ov = openpyxl.load_workbook(a.new, data_only=True), openpyxl.load_workbook(a.old, data_only=True)
    nf, of = openpyxl.load_workbook(a.new), openpyxl.load_workbook(a.old)
    diff = json.loads(a.tkdiff.read_text(encoding="utf-8"))
    out = openpyxl.Workbook()
    B = Font(bold=True)
    summary = {}
    ws = out.active
    ws.title = "①TK名稱新舊值"
    ws.append(["名稱", "標籤", "舊值（v19.1）", "新值（v21.2）", "變動率", "公式引用處（頁：格）", "本模型用法"])
    refs = {}
    for sh in nf.worksheets:
        if sh.title in ("TK_Link", "_Defaults"):
            continue
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for m in set(re.findall(r"\bTK_[A-Za-z0-9_]+", c.value)):
                        refs.setdefault(m, []).append(f"{sh.title}!{c.coordinate}")
    for ch in diff["changed"]:
        c0 = ch["cols"][0]
        r = refs.get("TK_" + ch["name"], [])
        use = "對照列／診斷（不回饋）" if r else "僅快照"
        ws.append([ch["name"], ch["label"], c0["old"], c0["new"], c0["rel"], "；".join(r), use])
    summary["tk_changed"] = len(diff["changed"])
    w2 = out.create_sheet("②關鍵輸出")
    w2.append(["輸出", "版本", "2025", "2026", "2027", "2028", "2029", "2030"])
    maxd = 0
    for nm, lab in KEY_OUT:
        o, n = named(ov, nm), named(nv, nm)
        w2.append([lab, "v0.1.4"] + o)
        w2.append([lab, "v0.1.5"] + n)
        d = [(y - x) if num(x) and num(y) else None for x, y in zip(o, n)]
        w2.append([lab, "差"] + d)
        maxd = max([maxd] + [abs(x) for x in d if x is not None])
    w2.append([])
    w2.append(["標量", "v0.1.4", "v0.1.5", "差"])
    for nm, lab in SCALARS:
        o, n = named(ov, nm)[0], named(nv, nm)[0]
        w2.append([lab, o, n, (n - o) if num(o) and num(n) else None])
        maxd = max(maxd, abs(n - o) if num(o) and num(n) else 0)
    summary["key_max_abs_diff"] = maxd
    w3 = out.create_sheet("③逐格比較")
    w3.append(["頁", "格", "類別", "舊值", "新值", "說明"])
    cats = {}
    tot = 0
    for sh in nv.sheetnames:
        x, y, fx, fy = ov[sh], nv[sh], of[sh], nf[sh]
        for r in range(1, max(x.max_row, y.max_row) + 1):
            for c in range(1, max(x.max_column, y.max_column) + 1):
                tot += 1
                v1, v2, f1, f2 = x.cell(r, c).value, y.cell(r, c).value, fx.cell(r, c).value, fy.cell(r, c).value
                if same(v1, v2) and f1 == f2:
                    continue
                coord = y.cell(r, c).coordinate
                if isinstance(v1, str) and isinstance(v2, str):
                    cat = "說明文字／版本字串"
                elif sh == "TK_Link":
                    cat = "TK_Link 快照值"
                elif sh in ("Cost", "Bridge", "EtaDiag", "Checks"):
                    cat = "對照列／診斷／資訊列（引用 TK 10 名稱）"
                else:
                    cat = "範圍外"
                cats[cat] = cats.get(cat, 0) + 1
                w3.append([sh, coord, cat, v1, v2, ""])
    w3.append([])
    w3.append(["比較格數（所有頁）", tot])
    for k, v in cats.items():
        w3.append([k, v])
    summary["cells_total"], summary["cats"] = tot, cats
    for w in out.worksheets:
        for c in w[1]:
            c.font = B
        w.column_dimensions["A"].width = 40
    out.save(a.out)
    if a.json:
        a.json.write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    print("saved", a.out, summary)


if __name__ == "__main__":
    main()
