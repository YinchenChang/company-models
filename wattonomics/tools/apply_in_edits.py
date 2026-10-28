# 把 In 頁輸入更正表（edits/v<版>_In_edits.csv）寫進現行 Excel 的 In 頁，產生新版的 --base 檔（Excel 優先：輸入只在 Excel 改）。
# 用法：python3 wattonomics/tools/apply_in_edits.py <現行.xlsx> <edits.csv> <輸出 base.xlsx>
# 欄 col 為 In 頁欄字母；old 非空時須與現值相等（數值比對容許 1e-9），否則停止。
import csv, sys
import openpyxl
from openpyxl.utils import column_index_from_string


def same(a, b):
    if b == "":
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a) == b


def num(v):
    try:
        f = float(v); return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def main(src, edits, out):
    wb = openpyxl.load_workbook(src)
    ws = wb["In"]
    rowof = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    n = 0
    for e in csv.DictReader(open(edits, encoding="utf-8-sig")):
        r = rowof[e["key"]]; c = column_index_from_string(e["col"])
        cur = ws.cell(r, c).value
        if not same(cur, e["old"]):
            sys.exit(f"停止：{e['key']} {e['col']}{r} 現值 {cur!r} ≠ old {e['old']!r}")
        ws.cell(r, c).value = num(e["new"]) if e["col"] not in ("E", "G") else e["new"]
        print(f"{e['key']} {e['col']}{r}: {cur!r} → {ws.cell(r, c).value!r}"); n += 1
    wb.save(out)
    print(f"{out}：{n} 格")


if __name__ == "__main__":
    main(*sys.argv[1:4])
