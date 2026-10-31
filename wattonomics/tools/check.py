# GW 帳核對程式：重算後活頁簿的錯誤格 0、Checks 全過（WGOV_Errors＝0）、In 頁標記與來源齊全、具名範圍可解析、快照一致（給快照時）。
# 用法：python3 wattonomics/tools/check.py <重算後的 Wattonomics xlsx> [<snapshot.json>]
# 任何 ERROR 以代碼 1 結束；WARN 只列出。
import json, os, sys
import openpyxl

ERRS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")
REQUIRED = ["README", "In", "Players", "CapexGW", "Funding", "Gap", "FundRaw", "Facility", "SegIn", "Spend", "Matrix", "Pipeline", "Deploy", "C0", "C1", "Compare", "TK_In", "SIL_In", "Interface", "Checks",
            "Decisions", "DB_Evidence", "Sources"]
TAGS = ("Verified", "Interested-party", "Analogy", "Assumed", "Derived", "待證據")


def main():
    x = sys.argv[1]
    wb = openpyxl.load_workbook(x, data_only=True)
    wf = openpyxl.load_workbook(x, data_only=False)
    bad = []
    # 1) 頁面
    miss = [s for s in REQUIRED if s not in wb.sheetnames]
    if miss: bad.append(f"缺頁面：{miss}")
    # 2) 錯誤格與未重算（公式格無快取值）
    nerr = 0; nuncalc = 0
    for ws in wb.worksheets:
        wsf = wf[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith(ERRS):
                    nerr += 1
                    if nerr <= 10: print(f"  錯誤格 {ws.title}!{c.coordinate}：{v}")
                fv = wsf[c.coordinate].value
                if isinstance(fv, str) and fv.startswith("=") and v is None:
                    nuncalc += 1
    if nerr: bad.append(f"錯誤格 {nerr}")
    if nuncalc > 0:
        print(f"  註：{nuncalc} 個公式格值為空（空字串結果或未重算；若全部為空表示未經 LibreOffice 重算）")
    # 3) Checks
    ck = wb["Checks"]
    rows = [(ck.cell(r, 1).value, ck.cell(r, 2).value, ck.cell(r, 6).value) for r in range(4, ck.max_row + 1) if ck.cell(r, 1).value]
    ce = [x_ for x_ in rows if x_[2] == "ERROR"]; cw = [x_ for x_ in rows if x_[2] == "WARN"]
    nok = sum(1 for x_ in rows if x_[2] == "OK")
    for i, l, s in ce: print(f"  Checks ERROR {i}：{l}")
    for i, l, s in cw: print(f"  Checks WARN  {i}：{l}")
    names = {}
    for nm, dn in wb.defined_names.items():
        names[nm] = dn.attr_text
    def val(nm):
        sh, ref = names[nm].split("!")
        return wb[sh.strip("'")][ref.replace("$", "")].value
    ge = val("WGOV_Errors")
    if ge != 0: bad.append(f"WGOV_Errors＝{ge}")
    if ce: bad.append(f"Checks ERROR {len(ce)}")
    # 4) In 頁標記與來源（獨立於 Excel 公式再檢一次）
    ws = wb["In"]; src_ids = {wb["Sources"].cell(r, 1).value for r in range(4, wb["Sources"].max_row + 1)}
    ntag = nsrc = nbad = 0; nrow = 0; pend = 0
    for r in range(5, ws.max_row + 1):
        if not ws.cell(r, 19).value: continue
        nrow += 1
        tag, sid = ws.cell(r, 5).value, ws.cell(r, 6).value
        if not tag: ntag += 1
        elif not any(t in str(tag) for t in TAGS): nbad += 1; print(f"  In!E{r} 標記不在清單：{tag}")
        if not sid: nsrc += 1
        elif sid not in src_ids: nbad += 1; print(f"  In!F{r} 來源 ID 不在 Sources：{sid}")
        if tag == "待證據": pend += 1
    if ntag or nsrc or nbad: bad.append(f"In 標記缺 {ntag}、來源缺 {nsrc}、無效 {nbad}")
    # 5) WIF_ 名稱
    gif = [n for n in names if n.startswith("WIF_")]
    for n in gif:
        v = val(n)
        if isinstance(v, str) and v.startswith(ERRS): bad.append(f"{n} 錯誤")
    # 6) 快照一致
    if len(sys.argv) > 2:
        snap = json.load(open(sys.argv[2], encoding="utf-8"))
        diff = [k for k, v in snap["names"].items() if k in names and val(k) != v]
        if diff: bad.append(f"快照與活頁簿不一致：{diff[:10]}")
        if snap["source"]["file"] != os.path.basename(x): bad.append("快照記錄的檔名不同")
    # 7) 固定規格（Andy 2026-10-28、2026-10-31）：README 須含對帳鏈、延遲情景與資金缺口規格
    rd = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "README.md")
    txt = open(rd, encoding="utf-8").read() if os.path.exists(rd) else ""
    for kw in ("### 對帳鏈", "### 基建延遲情景", "| C0 |", "| C7 |", "D0 準時", "D1 基準延遲", "D2 嚴重延遲", "IT 關鍵電力", "IF_FacilityGW", "k_chip2it", "下游取用 2027 年後數值前看覆蓋率", "Wattonomics（瓦帳）＝全市場 GW 唯一基準事實", "舊名 GW 帳／gwledger",
               # 資金缺口固定規格（Andy 2026-10-31 04:13、04:20、10:22）
               "資金缺口（Gap 頁", "新增 GW × 每 GW 建置成本", "不扣股票買回與股利", "整個市場當一間公司"):
        if kw not in txt: bad.append(f"wattonomics/README.md 缺固定規格：{kw}")
    print(f"{os.path.basename(x)}：頁面 {len(wb.sheetnames)}、具名範圍 {len(names)}（WIF_ {len(gif)}）、Checks {len(rows)} 項（OK {nok}／WARN {len(cw)}／ERROR {len(ce)}）、"
          f"In 輸入列 {nrow}（待證據 {pend}）、錯誤格 {nerr}")
    if bad:
        print("不通過：", "；".join(bad)); sys.exit(1)
    print("通過")


if __name__ == "__main__":
    main()
