# Wattonomics 互動頁取數：由重算後活頁簿匯出扁平 JSON（每筆一列，含標記與來源 ID）。
# 用法：python3 wattonomics/tools/export_explorer.py <重算後的 Wattonomics xlsx> <輸出 explorer.json>
# 內容：玩家 GW 年底（主路徑、SiteOnly）、目標 GW 與時點（In T2 園區管線）、資本支出（In Y2、CapexGW 甲2）、資金來源（Funding）、
#       Spend 拆分、Facility 每 GW 拆分、覆蓋率、市場總 GW（A×D0／D1／D2）。數值只讀 Interface 的 WIF_ 名稱或其來源頁（不重算）。
import hashlib, json, os, sys
import openpyxl


def main():
    x, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(x, data_only=True)
    names = {nm: dn.attr_text for nm, dn in wb.defined_names.items()}

    def val(nm):
        sh, ref = names[nm].split("!")
        v = wb[sh.strip("'")][ref.replace("$", "")].value
        return None if v == "" else v

    IN = wb["In"]
    meta = {}
    for r in range(5, IN.max_row + 1):
        k = IN.cell(r, 1).value
        if k and IN.cell(r, 19).value:
            meta[k] = dict(row=r, label=IN.cell(r, 2).value, basis=IN.cell(r, 3).value, unit=IN.cell(r, 4).value,
                           tag=IN.cell(r, 5).value, src=IN.cell(r, 6).value, note=IN.cell(r, 7).value)
    years = list(range(2024, 2033))
    recs = []

    def add(kind, **kw):
        kw["kind"] = kind; recs.append(kw)

    # 1) 市場總 GW 與覆蓋率
    for path, sfx in (("main", ""), ("site_only", "_SiteOnly")):
        for s in "ABC":
            for d in ("D0", "D1", "D2"):
                for y in years:
                    nm = f"WIF_MktGW_YE_{s}_{d}{sfx}_{y}"
                    if nm in names:
                        add("market_gw", path=path, scen=s, delay=d, year=y, value=val(nm), unit="IT GW", tag="Derived", src="Deploy", name=nm)
    for y in years:
        add("coverage", year=y, value=val(f"WIF_Coverage_{y}"), unit="比", tag="Derived", src="Deploy", name=f"WIF_Coverage_{y}",
            note="有證據支撐的年新增 ÷ 主路徑年新增；下游取用 2027 年後數值前看此值")
    # 2) 玩家年底 GW（主路徑／SiteOnly），標記取 In Y1 輸入列
    pkeys = [k[3:] for k in meta if k.startswith("mw_")]
    for p in pkeys:
        m = meta[f"mw_{p}"]
        for path, sfx in (("main", ""), ("site_only", "_SiteOnly")):
            for y in years:
                nm = f"WIF_PlayerGW_YE_{p}{sfx}_{y}"
                if nm in names:
                    add("player_gw", player=p, player_name=m["label"].replace(" 年底上線 MW", ""), path=path, year=y, value=val(nm), unit="IT GW",
                        tag=("Derived（甲2）" if path == "main" and y >= 2027 and p in ("msft", "googl", "amzn", "meta", "orcl", "crwv", "nbis") else m["tag"]),
                        src=m["src"], name=nm)
    # 3) 目標 GW 與時點（園區管線 T2）
    for k, m in meta.items():
        if not k.startswith(("cp_", "cpx_")): continue
        r = m["row"]
        g = lambda c: IN.cell(r, c).value
        add("campus", key=k, name_=m["label"], operator=g(8), user=g(9), mw=g(10), basis=m["basis"], online_year=g(11), online_month=g(12),
            known_delay_m=g(13), included=g(14), power_secured=g(15), tag=m["tag"], src=m["src"], note=m["note"])
    # 4) 資本支出：In Y2（個股模型，財年）與 CapexGW 甲2
    for k, m in meta.items():
        if not k.startswith("cap_"): continue
        r = m["row"]
        for j, y in enumerate(range(2023, 2033)):
            v = IN.cell(r, 8 + j).value
            if v not in (None, ""):
                add("capex_model", player=k[4:], fiscal_year=y, value=v, unit="US$B", tag=m["tag"], src=m["src"], note=m["note"])
    CX = wb["CapexGW"]
    for r in range(1, CX.max_row + 1):
        key = CX.cell(r, 2).value
        if isinstance(key, str) and (key.endswith("_gw") or key.endswith("_cy") or key.endswith("_ai")):
            p, f = key.rsplit("_", 1)
            for j, y in enumerate(range(2023, 2033)):
                v = CX.cell(r, 4 + j).value
                if isinstance(v, (int, float)):
                    add({"gw": "capex_gw", "cy": "capex_calendar", "ai": "capex_ai"}[f], player=p, year=y, value=v,
                        unit="IT GW" if f == "gw" else "US$B", tag="Derived", src="CapexGW", note="輸入為 Interested-party／共識（In Y2）")
    # 5) 資金來源（Funding 頁每列）
    FD = wb["Funding"]
    hdr = [FD.cell(4, c).value for c in range(1, FD.max_column + 1)]
    for r in range(5, FD.max_row + 1):
        if FD.cell(r, 2).value is None or not isinstance(FD.cell(r, 3).value, int): break
        row = {hdr[c - 1]: FD.cell(r, c).value for c in range(1, len(hdr) + 1) if hdr[c - 1]}
        add("funding", player=row["鍵"], player_name=row["玩家"], year=row["期間"], period=row["期間標籤"],
            values={k: (None if v == "" else v) for k, v in row.items() if k not in ("玩家", "鍵", "期間", "期間標籤")},
            unit="USD B（% 欄為比例）", tag="Derived（輸入見 FundRaw 標記）", src="SRC_GL_300／SRC_GL_302",
            default_view="gross",
            gross={"external": row.get("【總額】外部資金"), "external_pct": row.get("【總額】外部 ÷ 資本支出"), "ocf_pct": row.get("【總額】營業現金流 ÷ 資本支出"),
                   "debt_pct": row.get("【總額】債 ÷ 資本支出"), "equity_pct": row.get("【總額】股 ÷ 資本支出"), "lease_jv_pct": row.get("【總額】租賃＋表外 ÷ 資本支出")},
            payment_order={"internal_pct": row.get("內部現金 %"), "lease_pct": row.get("融資租賃 %"), "equity_pct": row.get("股權 %"), "debt_pct": row.get("債 %"), "other_pct": row.get("其餘 %")},
            note="gross＝期間外部資金全額 ÷ 資本支出（預設顯示）；payment_order＝支付順序（外部資金排在營業現金流之後）")
    FR = wb["FundRaw"]
    for r in range(4, FR.max_row + 1):
        if FR.cell(r, 1).value is None: continue
        g = lambda c: FR.cell(r, c).value
        add("funding_raw", id=g(1), player=g(3) or g(2), period=g(4), bucket=g(5), item=g(6), value=g(7), unit=g(8), detail=g(9), tag=g(10),
            source=g(11), url=g(12), verdict=g(13), correction=g(14), used=g(15), src="SRC_GL_300")
    # 6) Spend 與 Facility
    for nm in names:
        if nm.startswith(("WIF_Spend_", "WIF_Facility_")):
            base, y = nm.rsplit("_", 1)
            if y.isdigit():
                add("spend" if nm.startswith("WIF_Spend_") else "facility", item=base.split("_", 2)[2], year=int(y), value=val(nm),
                    unit=("US$B" if nm.startswith("WIF_Spend_") else ("GW" if "_mkt_q_" in nm else "$B/IT GW")), tag="Derived",
                    src="Spend／Facility（SRC_GL_301、SRC_GL_303）", name=nm)
    # 7) 資金缺口（v0.1.7；kind＝gap，part 區分：market／verdict／pergw／player／total）。只讀 WIF_Gap_* 名稱
    GY = list(range(2024, 2031))
    MF = [("need_rest", "NeedRest"), ("mkt_add", "MktAdd")]
    for vx, nx in (("", ""), ("_mc", "_ModelCapex")):
        MF += [("need" + vx, "Need" + nx), ("need_own" + vx, "NeedOwn" + nx), ("short" + vx, "Short" + nx), ("short_incl_rest" + vx, "ShortInclRest" + nx),
               ("cum_need" + vx, "CumNeed" + nx), ("cum_need_incl_rest" + vx, "CumNeedInclRest" + nx),
               ("cum_unarranged" + vx, "CumUnarranged" + nx), ("cum_unarranged_incl_rest" + vx, "CumUnarrangedInclRest" + nx),
               ("cum_memo_incl_committed" + vx, f"CumUnarranged{nx}_MemoInclCommitted"), ("cum_memo_incl_committed_incl_rest" + vx, f"CumUnarrangedInclRest{nx}_MemoInclCommitted")]
    if any(n.startswith("WIF_Gap_") for n in names):
        for s in "ABC":
            for d in ("D0", "D1", "D2"):
                for y in GY:
                    v = {k: (val(f"WIF_Gap_{nm}_{s}_{d}_{y}") if f"WIF_Gap_{nm}_{s}_{d}_{y}" in names else None) for k, nm in MF}
                    v.update(nonai=val(f"WIF_Gap_NonAI_{y}"), ocf=val(f"WIF_Gap_OCF_{y}"), buyback=val(f"WIF_Gap_Buyback_{y}"),
                             fallback=val(f"WIF_Gap_Fallback_{y}"), fallback_cells=val(f"WIF_Gap_FallbackCells_{y}"),
                             ext_obs=(val(f"WIF_Gap_ExtObs_{y}") if f"WIF_Gap_ExtObs_{y}" in names else None))
                    add("gap", part="market", scen=s, delay=d, year=y, values=v, unit="US$B", tag="Derived", src="Gap",
                        name=f"WIF_Gap_*_{s}_{d}_{y}", note="需要＝11 家 AI 資料中心＋市場其餘（資金來源未知）＋11 家非 AI；來源＝11 家營業現金流（個股模型基準，不隨情景變）＋已觀測外部融資（2024–2025）")
                add("gap", part="verdict", scen=s, delay=d, verdict=val(f"WIF_Gap_Verdict_{s}_{d}"), verdict12=val(f"WIF_Gap_Verdict12_{s}_{d}"),
                    verdict_mc=val(f"WIF_Gap_Verdict_ModelCapex_{s}_{d}"), verdict12_mc=val(f"WIF_Gap_Verdict12_ModelCapex_{s}_{d}"), tag="Derived", src="Gap",
                    note="G3：2030 自由現金流 < 0 或缺口占資本支出 > 30% 連 3 年＝OpenAI 型；否則 Anthropic 型")
        for nm in names:
            if nm.startswith("WIF_Gap_PerGW_") or nm.startswith("WIF_Gap_Item_"):
                base, y = nm.rsplit("_", 1)
                item = "total" if nm.startswith("WIF_Gap_PerGW_") else base[len("WIF_Gap_Item_"):]
                add("gap", part="pergw", item=item, year=int(y), value=val(nm), unit="$B/IT GW", tag="Derived", src="Gap（TK IF_CapexTotal）", name=nm)
        GPS = sorted({nm.split("_")[3] for nm in names if nm.startswith("WIF_Gap_Player_")})
        for p in GPS:
            for y in GY:
                v = {k: val(f"WIF_Gap_Player_{p}_{y}_{k}") for k in ("OCF", "CapexAI", "CapexNonAI", "Gap", "SelfFund", "Buyback", "CapexAIModel", "CapexAIGW", "CapexAIMC") if f"WIF_Gap_Player_{p}_{y}_{k}" in names}
                add("gap", part="player", player=p, year=y, values=v, unit="US$B（SelfFund 為倍數）", tag="Derived", src="Gap", name=f"WIF_Gap_Player_{p}_{y}_*")
            add("gap", part="player_meta", player=p, open_cash=val(f"WIF_Gap_PlayerCash_{p}"),
                committed=(val(f"WIF_Gap_PlayerCommitted_{p}") if f"WIF_Gap_PlayerCommitted_{p}" in names else None),
                verdict=val(f"WIF_Gap_PlayerVerdict_{p}"), tag="Derived", src="Gap")
        add("gap", part="total", open_cash=val("WIF_Gap_OpenCash"), committed=val("WIF_Gap_Committed"), ext2026_memo=val("WIF_Gap_ExtObs2026Memo"),
            buffer_cum=val("WIF_Gap_BufferCum"), ocf_missing={y: val(f"WIF_Gap_OCFMissing_{y}") for y in GY}, coverage={y: val(f"WIF_Coverage_{y}") for y in GY},
            tag="Derived", src="Gap", note="期初現金只進累計；已簽未起租與表外合資承諾是市場內移轉，只列備忘（G4 更正）；2026 年已觀測融資只列備忘；買回股利不扣（Andy 2026-10-31 04:20）")
    snap = dict(source=dict(file=os.path.basename(x), sha256=hashlib.sha256(open(x, "rb").read()).hexdigest(),
                            note="Wattonomics（瓦帳）互動頁取數；扁平紀錄，kind 欄區分種類；單位 IT GW（＝TK GW）；下游取用 2027 年後數值前看 coverage",
                            wgov_errors=val("WGOV_Errors")),
                kinds=sorted({r["kind"] for r in recs}), records=recs)
    json.dump(snap, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    from collections import Counter
    print(f"explorer：{out}（{len(recs)} 筆；{dict(Counter(r['kind'] for r in recs))}）")


if __name__ == "__main__":
    main()
