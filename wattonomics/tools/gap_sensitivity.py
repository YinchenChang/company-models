# Wattonomics Gap 頁敏感度：把 In 頁某純量的基準值（I 欄）改成低／高（H／J 欄），LibreOffice 重算後讀 WIF_Gap_* 累計值與判讀。
# 用法：python3 wattonomics/tools/gap_sensitivity.py <重算後的 Wattonomics xlsx> [<輸出 json>]
# 只在暫存目錄產生變體檔，不改原檔；所有數字仍由 Excel 公式算出（本程式只改輸入格、讀輸出名稱）。
import json, os, shutil, subprocess, sys, tempfile
import openpyxl

VARIANTS = [("基準", {}), ("lag_bldg 低", {"lag_bldg": "lo"}), ("lag_bldg 高", {"lag_bldg": "hi"}),
            ("ai_sh 低（含 ai_amzn）", {"ai_sh": "lo", "ai_amzn": "lo"}), ("ai_sh 高（含 ai_amzn）", {"ai_sh": "hi", "ai_amzn": "hi"})]
SCENS = [("A", "D0"), ("A", "D1"), ("A", "D2")]
OUT = ("CumNeed", "CumUnarranged", "CumNeed_ModelCapex", "CumUnarranged_ModelCapex")


def read(x):
    wb = openpyxl.load_workbook(x, data_only=True)
    def v(nm):
        sh, ref = wb.defined_names[nm].attr_text.split("!")
        return wb[sh.strip("'")][ref.replace("$", "")].value
    res = {}
    for s, d in SCENS:
        k = f"{s}_{d}"
        res[k] = {o: {y: v(f"WIF_Gap_{o}_{k}_{y}") for y in range(2027, 2031)} for o in OUT}
        res[k]["Verdict"] = v(f"WIF_Gap_Verdict_{k}"); res[k]["VerdictMC"] = v(f"WIF_Gap_Verdict_ModelCapex_{k}")
    return res


def main():
    src = sys.argv[1]; out = sys.argv[2] if len(sys.argv) > 2 else None
    tmp = tempfile.mkdtemp(prefix="gapsens_")
    allres = {}
    for lab, ch in VARIANTS:
        if not ch:
            allres[lab] = read(src); continue
        wb = openpyxl.load_workbook(src)
        ws = wb["In"]
        for key, side in ch.items():
            sh, ref = wb.defined_names[key].attr_text.split("!")
            row = int(ref.replace("$", "")[1:])
            col = {"lo": "H", "hi": "J"}[side]
            ws[f"I{row}"] = ws[f"{col}{row}"].value
        fn = os.path.join(tmp, f"v{len(allres)}.xlsx"); wb.save(fn)
        od = os.path.join(tmp, f"rc{len(allres)}"); os.makedirs(od)
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", od, fn], check=True, capture_output=True)
        allres[lab] = read(os.path.join(od, os.path.basename(fn)))
    shutil.rmtree(tmp, ignore_errors=True)
    for lab, r in allres.items():
        print(lab)
        for k, d in r.items():
            print(f"  {k}: 累計尚需 2030 {d['CumNeed'][2030]:.0f}；仍未安排 {d['CumUnarranged'][2030]:.0f}；並列讀法 {d['CumNeed_ModelCapex'][2030]:.0f}／{d['CumUnarranged_ModelCapex'][2030]:.0f}；判讀 {d['Verdict']}｜並列 {d['VerdictMC']}")
    if out:
        json.dump(allres, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
