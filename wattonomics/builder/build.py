# Wattonomics（瓦帳；舊名 瓦帳）建檔程式
# 用法：
#   首版（由種子建）：python3 wattonomics/builder/build.py --out build/raw/<YYYYMMDD>_Wattonomics_v0.1.xlsx
#   之後各版（Excel 優先）：python3 wattonomics/builder/build.py --base wattonomics/model/<現行>.xlsx --out <raw.xlsx>
#     --base 時，In 頁的藍字輸入（口徑、標記、來源、說明、各年值或低／基準／高）一律讀回現行 Excel；種子只用於 Excel 裡還沒有的鍵。
# 產出後以 LibreOffice headless 重算存檔（見 wattonomics/README.md），重算錯誤須為 0。
# 原則（比照 siliconomics/builder/build.py）：Excel 是唯一事實來源；本程式只產生公式頁與格式，計算一律為 Excel 公式。
import argparse, json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
GL = os.path.dirname(HERE)
ROOT = os.path.dirname(GL)
VERSION = "v0.1.7"
DATE = "20261031"
TK_SNAPSHOT = os.path.join(ROOT, "siliconomics", "data", "tokenomics_snapshot_v21.7.json")
SIL_DIR = os.path.join(ROOT, "siliconomics", "data")
SEED = os.path.join(HERE, "seed", "seed_v0.1.json")
TK_EXTRA = os.path.join(HERE, "seed", "tk_extra_v21.7.json")   # 快照未含、直接讀 Tokenomics 活頁簿的名稱（R1 裁定）
EVID = os.path.join(HERE, "seed", "evidence_v0.1.json")         # tools/import_evidence.py 由 evidence_raw 產生
FUND = os.path.join(HERE, "seed", "funding_v0.1.4.json")        # tools/import_funding.py 由 F_funding.csv、P_facility_power.csv 產生（附 G 段判定）
CAMPUS_ADD = sorted(__import__("glob").glob(os.path.join(HERE, "seed", "campus_v*_add.json")))   # v0.1.2 起新增的 T2 園區列（各版一檔；園區候選回查＋獨立抽查後合併；值之後由 Excel 讀回）
SEG = os.path.join(ROOT, "markets", "invest_map", "data", "segments_ai_semis_v0.2.3.csv")   # AI半導體 次產業比例（最新版）
SEG_COMMIT = "5dea5c8"
GAPSEED = os.path.join(HERE, "seed", "gap_models_v0.1.7.json")   # tools/import_gap_models.py：Gap 頁個股模型基準值（只用於 In Y9 首建）
GP = ["msft", "googl", "amzn", "meta", "orcl", "crwv", "nbis", "wyfi", "xai", "oai", "ant", "zhipu"]   # Gap 頁玩家（智譜只列，不進合計）
GPN = {"msft": "Microsoft", "googl": "Alphabet", "amzn": "Amazon", "meta": "Meta", "orcl": "Oracle", "crwv": "CoreWeave", "nbis": "Nebius",
       "wyfi": "WhiteFiber", "xai": "SpaceX／xAI", "oai": "OpenAI", "ant": "Anthropic", "zhipu": "智譜"}
GSRC = {p: f"SRC_GL_{310 + i}" for i, p in enumerate(GP)}

BLUE = Font(color="0000FF"); GREEN = Font(color="008000"); BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7"); KEY = PatternFill("solid", fgColor="FFF2CC")
PEND = PatternFill("solid", fgColor="FFFF00"); WRAP = Alignment(wrap_text=True, vertical="top")
YEARS = list(range(2023, 2033))           # 2023＝期初存量年；輸出 2024–2032
OUTY = [y for y in YEARS if y >= 2024]
YC = {y: get_column_letter(4 + i) for i, y in enumerate(YEARS)}   # 計算頁：D..M
IC = {y: get_column_letter(8 + i) for i, y in enumerate(YEARS)}   # In 頁：H..Q
SCEN = ["A", "B", "C"]
DLY = ["D0", "D1", "D2"]
GENS = ["H100", "GB200", "GB300", "VR200", "RU"]
GEN_NAME = {"H100": "Hopper H100", "GB200": "GB200 NVL72", "GB300": "GB300 NVL72", "VR200": "VR200 NVL72", "RU": "Rubin Ultra"}


def hdr(ws, row, labels, c0=1):
    for j, h in enumerate(labels, c0):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def load_base_inputs(path):
    """讀回現行 Excel 的 In 頁（Excel 優先）：{鍵: {欄: 值}}（C、E–Q 欄）"""
    if not path:
        return {}
    ws = openpyxl.load_workbook(path, data_only=False)["In"]
    out = {}
    for r in range(5, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if not k or not ws.cell(r, 19).value:
            continue
        out[k] = {c: ws.cell(r, c).value for c in [3] + list(range(5, 18))}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base")
    ap.add_argument("--out", required=True)
    ap.add_argument("--version")
    ap.add_argument("--date")
    a = ap.parse_args()
    if a.base:
        fn = os.path.basename(a.base)
        if a.version is None and "_v" in fn:
            a.version = "v" + fn.rsplit("_v", 1)[1].rsplit(".xlsx", 1)[0]
        if a.date is None and fn[:8].isdigit():
            a.date = fn[:8]
    a.version = a.version or VERSION
    a.date = a.date or DATE
    seed = json.load(open(SEED, encoding="utf-8"))
    for fp in CAMPUS_ADD:
        add = json.load(open(fp, encoding="utf-8"))
        have = {c["key"] for c in seed["campus"]}
        assert not have & {c["key"] for c in add["campus"]}, "新增園區鍵與既有重複"
        seed["campus"] += add["campus"]
        seed["sources"] += [x for x in add.get("sources", []) if x["id"] not in {y["id"] for y in seed["sources"]}]
    base = load_base_inputs(a.base)
    tk = json.load(open(TK_SNAPSHOT, encoding="utf-8"))
    tkx = json.load(open(TK_EXTRA, encoding="utf-8"))
    sil_fn = open(os.path.join(SIL_DIR, "CURRENT_SNAPSHOT"), encoding="utf-8").read().strip()
    sil = json.load(open(os.path.join(SIL_DIR, sil_fn), encoding="utf-8"))
    P = seed["players"]; PK = [p["key"] for p in P]; PN = {p["key"]: p["name"] for p in P}
    U = seed["users"]

    wb = openpyxl.Workbook()
    RD = wb.active; RD.title = "README"
    IN = wb.create_sheet("In")
    TKS = wb.create_sheet("TK_In")
    SL = wb.create_sheet("SIL_In")
    PL = wb.create_sheet("Players")
    MX = wb.create_sheet("Matrix")
    PP = wb.create_sheet("Pipeline")
    CX = wb.create_sheet("CapexGW")
    DP = wb.create_sheet("Deploy")
    C0 = wb.create_sheet("C0")
    C1 = wb.create_sheet("C1")
    CP = wb.create_sheet("Compare")
    FD = wb.create_sheet("Funding")
    GA = wb.create_sheet("Gap")
    FR = wb.create_sheet("FundRaw")
    FC = wb.create_sheet("Facility")
    SG = wb.create_sheet("SegIn")
    SP = wb.create_sheet("Spend")
    IF = wb.create_sheet("Interface")
    CK = wb.create_sheet("Checks")
    DC = wb.create_sheet("Decisions")
    EV = wb.create_sheet("DB_Evidence")
    SR = wb.create_sheet("Sources")
    RQ = wb.create_sheet("Requests")

    # ================= In：藍字輸入（Excel 擁有） =================
    IN["A1"] = ("In — 瓦帳輸入（藍字＝輸入，由本活頁簿擁有；要改值直接改本頁，builder 重建時以 --base 讀回）。"
                "純量列：H／I／J＝低／基準／高，具名範圍＝鍵（I 欄）；年度列：H–Q＝2023–2032；T1 分配表、T2 園區表見節標題")
    IN["A2"] = ("標記：Verified／Interested-party（寫明是誰）／Analogy／Assumed（給低／高）／Derived（寫明換算）／待證據（值留空）。"
                "來源 ID 見 Sources 頁；說明欄的 A#n／B#n／D#n＝evidence_raw 列號（DB_Evidence 頁，含 C 段判定）。"
                "口徑：IT／設施／未說明／晶片（主口徑 IT GW；設施 ÷ IF_FacilityGW；未說明者基準取 IT 讀法）。S 欄＝輸入列類型（治理檢查用）")
    hdr(IN, 4, ["鍵", "項目", "口徑", "單位", "標記", "來源 ID", "說明"] + [str(y) for y in YEARS] + ["", "類型"])
    INROW = {}; INKIND = {}
    r = 5

    def B(key, col, default):
        if key in base and col in base[key]:
            return base[key][col]
        return default

    def sec(title):
        nonlocal r
        IN.cell(r, 1, title).font = BOLD; r += 1

    def meta(key, label, basis, unit, tag, src, note, kind):
        IN.cell(r, 1, key); IN.cell(r, 2, label)
        IN.cell(r, 3, B(key, 3, basis)); IN.cell(r, 4, unit)
        IN.cell(r, 5, B(key, 5, tag)); IN.cell(r, 6, B(key, 6, src)); IN.cell(r, 7, B(key, 7, note))
        IN.cell(r, 19, kind)
        if B(key, 5, tag) == "待證據":
            for j in range(1, 8): IN.cell(r, j).fill = PEND
        INROW[key] = r; INKIND[key] = kind

    def scalar(s):
        nonlocal r
        k = s["key"]
        meta(k, s["label"], s.get("basis", "—"), s["unit"], s["tag"], s["src"], s.get("note", ""), 1)
        for col, v in ((8, s.get("lo")), (9, s.get("base")), (10, s.get("hi"))):
            v = B(k, col, v)
            if v is not None:
                c = IN.cell(r, col, v); c.font = BLUE
                if s["unit"] == "%": c.number_format = "0.0%"
        name(wb, k, f"In!$I${r}")
        if B(k, 8, s.get("lo")) is not None: name(wb, k + "_lo", f"In!$H${r}")
        if B(k, 10, s.get("hi")) is not None: name(wb, k + "_hi", f"In!$J${r}")
        r += 1

    def yrow(key, label, basis, unit, tag, src, note, vals, kind=2, fmt="#,##0.0##"):
        nonlocal r
        meta(key, label, basis, unit, tag, src, note, kind)
        for y in YEARS:
            v = B(key, 8 + YEARS.index(y), vals.get(str(y)))
            if v is not None:
                c = IN.cell(r, 8 + YEARS.index(y), v); c.font = BLUE; c.number_format = fmt
        r += 1

    for s in seed["scalars"]:
        if "sec" in s:
            sec(s["sec"]); continue
        scalar(s)
    sec("S5. 玩家非 NVIDIA 占比覆寫（未列者＝市場 NonNV 比例）")
    for p, o in seed["nnv"].items():
        if p.startswith("_"): continue
        scalar(dict(key=f"nnv_{p}", label=f"{PN[p]} 新增 GW 非 NVIDIA 占比", unit="%", base=o["base"], tag=o["tag"], src=o["src"], note=o["note"]))
    sec("S6. 個股模型財年結束月（C1 曆年→財年換算）")
    FYE = {}
    for p, o in seed["capex"].items():
        if p.startswith("_"): continue
        fsrc = seed.get("fye_src", {}).get(p, o["src"])
        scalar(dict(key=f"fye_{p}", label=f"{PN[p]} 財年結束月", unit="月", base=o["fye"], tag="Verified", src=fsrc,
                    note=("company.json calendar.fiscalYearEndMonth" if p != "meta" else "Meta 財年＝曆年（A#63 8-K Ex.99.1 期末 2026-09-30 為第三季）")))
        FYE[p] = f"fye_{p}"
    S_LAST = r - 1

    sec("Y0. Epoch AI 晶片功耗（MW，晶片層；IT＝× 倍數：Google、Amazon 用 C 段 2025 更正值 ÷ 2025 晶片功耗，其他用 k_chip2it）")
    CHIP = {}
    for p in PK:
        o = seed["chip"].get(p)
        if o:
            yrow(f"chip_{p}", f"{PN[p]} 擁有 AI 晶片功耗（年底）", "晶片", "MW", "Analogy", "SRC_GL_203" if p == "xai" else "SRC_GL_200", o["note"], o["y"], fmt="#,##0.0")
            CHIP[p] = INROW[f"chip_{p}"]
    sec("Y1. 玩家年底上線 MW（營運者口徑；輸入優先於晶片換算；空白＝由晶片、觀測、內插或園區延伸補值，見 Players 與 Decisions GL01-03）")
    for p in PK:
        o = seed["mw"].get(p)
        if o:
            yrow(f"mw_{p}", f"{PN[p]} 年底上線 MW", o["basis"], "MW", o["tag"], o["src"], o["note"], o["y"], fmt="#,##0.0")
        else:
            yrow(f"mw_{p}", f"{PN[p]} 年底上線 MW", "IT", "MW", "Derived", "SRC_GL_203", "無直接年底輸入：由 Y0 晶片 × 倍數、S3 觀測與 T2 園區延伸推得（Players）", {}, fmt="#,##0.0")
    sec("Y2. 個股模型資本支出（公司財年；欄＝財年，US$B；PR #190 未合併前讀現行 company.json／data）")
    for p in PK:
        o = seed["capex"].get(p)
        if o and o["y"]:
            yrow(f"cap_{p}", f"{PN[p]} 資本支出（財年）", "—", "US$B", o["tag"], o["src"], o["note"], o["y"], fmt="#,##0.0")
        else:
            yrow(f"cap_{p}", f"{PN[p]} 資本支出（財年）", "—", "US$B", "待證據", "SRC_GL_000", (o or {}).get("note", "待證據（無個股模型或未取得）"), {}, fmt="#,##0.0")
    sec("Y3. 非 AI 資本支出（公司財年，US$B；C1 加回）")
    for p in PK:
        yrow(f"nai_{p}", f"{PN[p]} 非 AI 資本支出（財年）", "—", "US$B", "待證據", "SRC_GL_000", "待證據；空白視為 0（C1 推算偏低）", {}, fmt="#,##0.0")
    sec("Y4. 使用者總量（報導或使用者模型；只作 Matrix 欄和對照）")
    for o in seed["usertotals"]:
        yrow(o["key"], o["label"], o["basis"], "GW", o["tag"], o["src"], o["note"], o["y"], fmt="0.00")
    sec("Y5. 第三方總量（C0 佐證；只作佐證，不校準）")
    for o in seed["thirdparty"]:
        yrow(o["key"], o["label"], o["basis"], "GW", o.get("tag", "待證據"), o.get("src", "SRC_GL_000"), o.get("note", "待證據"), o.get("y", {}), fmt="0.00")
    sec("Y6. 世代組合（新增 GW；NVIDIA 內部占比；空白年＝Players 頁由 SIL_In 推導〔Derived〕）")
    mx = seed["mix"]
    for g in GENS:
        yrow(f"mix_{g}", f"新增 GW 世代占比：{GEN_NAME[g]}", "—", "%", mx["tag_nv"], mx["src_nv"], mx["note_nv"], mx["nv"].get(g, {}), fmt="0%")
    yrow("nnv_mkt", "市場新增 GW 非 NVIDIA 占比（覆寫；空白＝SIL 推導，2023–25 沿用 2026、2030–32 沿用 2029）", "—", "%", "待證據", "SRC_GL_000",
         "2026–29 預設由 Siliconomics Units × 每顆 IT capex 推導（Derived）；2023–25 無來源，沿用 2026 值待決", {}, fmt="0%")
    # ---- Y9：資金缺口（Gap 頁）個股模型基準值；首建讀 seed，之後以本頁為準 ----
    GSEED = json.load(open(GAPSEED, encoding="utf-8")) if os.path.exists(GAPSEED) else {"players": {}}
    sec("Y9. 資金缺口：個股模型基準值（欄＝公司財年；Gap 頁依財年結束月換成曆年；tools/import_gap_models.py 首建，之後以本頁為準；空白＝無模型值，不外推、不補估）")
    GITEM = (("ocf", "營業現金流"), ("capai", "AI 資本支出（模型；對照）"), ("nai", "非 AI 資本支出"), ("cap", "資本支出合計（模型；對照）"), ("bb", "買回＋股利（備忘）"))
    for p in GP:
        g = GSEED["players"].get(p, {})
        src_ = f"{g.get('file', '')}（commit {g.get('commit', '')}）"
        for it, lab in GITEM:
            d = g.get("items", {}).get(it, {"values": {}, "cells": {}, "labels": [], "note": "無模型值"})
            vals = {k: v for k, v in d["values"].items() if v is not None}
            if vals:
                miss = [k for k, v in d["values"].items() if v is None]
                cells = "、".join(f"FY{k} {c}" for k, c in d["cells"].items() if d["values"].get(k) is not None)
                note = (f"{src_}；{cells}；標籤 {'；'.join(d['labels'])}；財年結束月 {g.get('fye', 12)}"
                        + (f"；無模型值：FY{'、FY'.join(miss)}" if miss else "") + (f"；{d['note']}" if d.get("note") else ""))
            else:
                note = f"無模型值（{src_}）：{d.get('note', '')}"
            yrow(f"g{it}_{p}", f"{GPN[p]} {lab}（財年，基準情境）", "—", g.get("unit", "US$B"), g.get("tag", "Derived"), GSRC[p], note, vals)
        c_ = g.get("cash", {})
        scalar(dict(key=f"gcash_{p}", label=f"{GPN[p]} 期初現金（{c_.get('asof', '')}）", unit=g.get("unit", "US$B"), base=c_.get("value"),
                    tag=g.get("tag", "Derived"), src=GSRC[p], note=f"{src_} {c_.get('cell', '')}「{c_.get('label', '')}」；{c_.get('asof', '')}"))
    fx_ = GSEED["players"].get("zhipu", {}).get("fx", {})
    scalar(dict(key="gfx_cny", label="智譜換算匯率 USD/CNY（元／美元；RMB 億 ÷ 10 ÷ 匯率＝US$B）", unit="元／美元", lo=fx_.get("lo"), base=fx_.get("value"), hi=fx_.get("hi"),
                tag="Verified", src=GSRC["zhipu"], note=f"智譜個股模型 {fx_.get('cell', '')}「{fx_.get('label', '')}」（{fx_.get('asof', '')}；中國外匯交易中心）；低／高＝模型 ±5%"))
    Y_LAST = r - 1

    UK = [u["key"] for u in U]
    sec("T1. 營運者 → 使用者分配（每列一筆：H 營運者鍵、I 使用者鍵〔" + "／".join(UK) + "〕、J 形式〔比＝占營運者年底 GW 比例；GW＝絕對值〕、K／L／M＝2025／2026／2027 年底；未列＝未分配）")
    hdr(IN, r, ["鍵", "項目", "口徑", "單位", "標記", "來源 ID", "說明", "營運者", "使用者", "形式", "2025", "2026", "2027"]); r += 1
    AL0 = r
    for al in seed["alloc"]:
        k = al["key"]
        meta(k, f"{PN[al['op']]} → {dict((u['key'], u['name']) for u in U)[al['user']]}", "IT", al["form"], al["tag"], al["src"], al["note"], 3)
        for j, v in enumerate((al["op"], al["user"], al["form"]) + tuple(al["y"])):
            v = B(k, 8 + j, v)
            if v is not None:
                c = IN.cell(r, 8 + j, v); c.font = BLUE
        r += 1
    AL1 = r - 1
    r += 1
    sec("T2. 園區管線（H 營運者鍵、I 使用者鍵、J MW、K 公告上線年、L 公告上線月、M 已知延遲〔月〕、N 納入〔1＝計入；2＝只進電力證據路徑，不入 Players／Pipeline〕、O 已確保電力〔1＝是；D2 不適用〕）")
    hdr(IN, r, ["鍵", "園區", "口徑", "單位", "標記", "來源 ID", "說明", "營運者", "使用者", "MW", "上線年", "上線月", "已知延遲月", "納入", "已確保電力"]); r += 1
    CAMP = []
    for cp in seed["campus"]:
        k = cp["key"]
        meta(k, cp["name"], cp["basis"], "MW", cp["tag"], cp["src"], cp["note"], 4)
        for j, f in enumerate(("op", "user", "mw", "year", "month", "delay", "inc", "sec")):
            v = B(k, 8 + j, cp.get(f))
            if v is not None and v != "":
                c = IN.cell(r, 8 + j, v); c.font = BLUE
        if not B(k, 14, cp.get("inc")):
            for j in range(1, 8): IN.cell(r, j).font = Font(color="808080")
        CAMP.append((k, r)); r += 1
    IN_LAST = r - 1
    for col, w in zip("ABCDEFGHIJKLMNOPQRS", (16, 50, 8, 7, 18, 12, 70) + (9,) * 11 + (5,)):
        IN.column_dimensions[col].width = w

    # ================= TK_In：Tokenomics 鏡像（只讀） =================
    src = tk["source"]
    TKS["A1"] = (f"TK_In — Tokenomics 取數（只讀鏡像；權威在 Tokenomics）：{src['file']}（{src['version']}，commit {src['commit'][:7]}，擷取 {src['extractedAt']}）。"
                 "同一份快照與 Siliconomics 相同（siliconomics/data/tokenomics_snapshot_v21.7.json），不另抽；綠字不得修改")
    TKS["A2"] = "快照註記：" + src.get("note", "")
    hdr(TKS, 4, ["名稱（TK）", "世代／項目", "本庫名稱", "低", "基準", "高", "單位", "TK 位置"])
    gens = {g["name"]: g["code"] for g in src["generations"]}
    rr = 5; TKN = set()
    allitems = list(tk["items"].items()) + [(k, dict(v, _extra=True)) for k, v in tkx["items"].items() if k not in tk["items"]]
    for nm_, it in allitems:
        short = nm_[3:]
        if it.get("_extra"):
            it = dict(it, ref=f"{tkx['source']['path']} {it['ref']}（快照未含；直接讀活頁簿，commit {tkx['source']['commit']}，SRC_GL_006）")
        if it.get("kind") == "gen_cost":
            for g, vv in it["values"].items():
                sname = f"TK_{short}_{gens.get(g, g)}"
                TKS.cell(rr, 1, nm_); TKS.cell(rr, 2, g); TKS.cell(rr, 3, sname)
                for j, c_ in ((4, "低成本"), (5, "基準"), (6, "高成本")):
                    TKS.cell(rr, j, vv[c_]).font = GREEN
                TKS.cell(rr, 7, it.get("unit")); TKS.cell(rr, 8, it.get("ref"))
                name(wb, sname, f"TK_In!$E${rr}"); name(wb, sname + "_lo", f"TK_In!$D${rr}"); name(wb, sname + "_hi", f"TK_In!$F${rr}")
                TKN.add(sname); rr += 1
        else:
            vals = it.get("values") or it.get("value")
            sname = f"TK_{short}"
            TKS.cell(rr, 1, nm_); TKS.cell(rr, 2, it.get("label")); TKS.cell(rr, 3, sname)
            if isinstance(vals, list):
                for j, x in zip((4, 5, 6), vals): TKS.cell(rr, j, x).font = GREEN
            else:
                TKS.cell(rr, 5, vals).font = GREEN
            TKS.cell(rr, 7, it.get("unit")); TKS.cell(rr, 8, it.get("ref"))
            name(wb, sname, f"TK_In!$E${rr}"); name(wb, sname + "_lo", f"TK_In!$D${rr}"); name(wb, sname + "_hi", f"TK_In!$F${rr}")
            TKN.add(sname); rr += 1
    TK_LAST = rr - 1
    HAS_FACGW = "IF_FacilityGW" in tk["items"]
    for col, w in zip("ABCDEFGH", (22, 40, 26, 10, 10, 10, 12, 20)):
        TKS.column_dimensions[col].width = w

    # ================= SIL_In：Siliconomics 鏡像（只讀） =================
    sn = sil["names"]; ssrc = sil["source"]
    SL["A1"] = (f"SIL_In — Siliconomics 取數（只讀鏡像；權威在 AI半導體）：{ssrc['file']}（{ssrc['version']}，sha256 {ssrc['sha256'][:12]}）。"
                "由 siliconomics/data/CURRENT_SNAPSHOT 指向；綠字不得修改")
    hdr(SL, 3, ["名稱（SIF_）", "本庫名稱", "單位", "值（純量）／2023"] + [str(y) for y in YEARS[1:]])
    SILS = {}; SILV = {}
    rr = 4
    SL.cell(rr, 1, "A. 純量").font = BOLD; rr += 1
    scal = [("SIF_vr26", "x"), ("SIF_vr26_lo", "x"), ("SIF_vr26_hi", "x")] + [(f"SIF_it_{k}", "US$K／顆") for k in ("GB", "VR", "RU", "HP", "AMD8", "AMDR", "TPU", "TRN", "ASIC")]
    for nm_, unit in scal:
        nn = "SIL_" + nm_[4:]
        SL.cell(rr, 1, nm_); SL.cell(rr, 2, nn); SL.cell(rr, 3, unit)
        SL.cell(rr, 4, sn[nm_]).font = GREEN
        name(wb, nn, f"SIL_In!$D${rr}"); SILS[nm_] = nn; rr += 1
    rr += 1
    SL.cell(rr, 1, "B. 年度（欄 D–M＝2023–2032；Siliconomics 只到 2029）").font = BOLD; rr += 1
    vec = []
    for k in ("bw", "rb", "ru", "hp", "amd", "tpu", "trn", "oth"):
        vec.append((f"SIF_UnitsPlat_{k}", 2025, "百萬顆"))
    for s in SCEN:
        vec.append((f"SIF_UnitsCoef_{s}", 2025, "x"))
    for s in SCEN:
        vec.append((f"SIF_DeployGW_{s}", 2026, "GW（成本等值）"))
    for s in SCEN:
        vec.append((f"SIF_AIITCapex_{s}", 2026, "US$B"))
    vec.append(("SIF_ITCapexPerGW", 2026, "US$B/GW"))
    for nm_, y0, unit in vec:
        SL.cell(rr, 1, nm_); SL.cell(rr, 3, unit)
        for i, v in enumerate(sn[nm_]):
            y = y0 + i
            if v is not None:
                SL.cell(rr, YEARS.index(y) + 4, v).font = GREEN
                SILV[(nm_, y)] = f"SIL_In!${YC[y]}${rr}"
        rr += 1
    for col, w in zip("ABCDEFGHIJKLM", (26, 16, 14, 14) + (10,) * 9):
        SL.column_dimensions[col].width = w

    def sv(nm_, y):
        return SILV[(nm_, y)]

    # ================= Players：玩家 × 年 =================
    PL["A1"] = ("Players — 玩家 × 年（2023–2032）年底上線 IT GW（＝TK GW；甲：公告路徑；營運者口徑）。已知值優先序：In Y1 輸入（設施 ÷ IF_FacilityGW）→ Y0 晶片功耗 × 倍數 → "
                "S3 觀測＋觀測後至年底的 T2 園區上線量；兩已知值間線性內插；最後已知值後＝前年＋當年 T2 園區（D0）；第一個已知值前留空（GL01-03）")
    PL["A2"] = "旗標：1 輸入／5 晶片換算／2 觀測／3 內插／4 園區延伸／0 未知。年新增＝年底差（第一個已知年為『新納入存量』，不計新增）。低／高列：晶片換算用 k_chip2it 區間或 C 段區間，其餘＝基準"
    hdr(PL, 3, ["項目", "鍵", "單位"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    PL.cell(r, 1, "世代組合（新增 GW；NVIDIA 內部占比）與換算係數").font = BOLD; r += 1
    MIXR = {}
    RUSH = r + len(GENS) + 1
    for g in GENS:
        ir = INROW[f"mix_{g}"]
        PL.cell(r, 1, f"世代占比：{GEN_NAME[g]}"); PL.cell(r, 2, f"mix_{g}"); PL.cell(r, 3, "%")
        for y in YEARS:
            c = YC[y]; inref = f"In!${IC[y]}${ir}"
            if y <= 2025:
                d = "0"
            elif y == 2026:
                d = {"GB300": "(1-SIL_vr26)", "VR200": "SIL_vr26"}.get(g, "0")
            elif y <= 2029:
                d = {"VR200": f"(1-{c}${RUSH})", "RU": f"{c}${RUSH}"}.get(g, "0")
            else:
                d = f"{YC[2029]}{r}"
            PL.cell(r, YEARS.index(y) + 4, f'=IF({inref}<>"",{inref},{d})').number_format = "0%"
        MIXR[g] = r; r += 1
    PL.cell(r, 1, "世代占比合計（須＝1）"); PL.cell(r, 3, "%")
    for y in YEARS:
        PL.cell(r, YEARS.index(y) + 4, f"=SUM({YC[y]}{MIXR[GENS[0]]}:{YC[y]}{MIXR[GENS[-1]]})").number_format = "0%"
    MIXSUM = r; r += 1
    assert r == RUSH
    PL.cell(r, 1, "Rubin Ultra 占 Rubin 平台 GW 比（Derived：Units ÷ TK 每 GW GPU 數）"); PL.cell(r, 3, "%")
    for y in (2027, 2028, 2029):
        ru, rb = sv("SIF_UnitsPlat_ru", y), sv("SIF_UnitsPlat_rb", y)
        PL.cell(r, YEARS.index(y) + 4, f"=({ru}/TK_GPUsPerGW_RU)/({rb}/TK_GPUsPerGW_VR200+{ru}/TK_GPUsPerGW_RU)").number_format = "0.0%"
    PL.cell(r, 14, "工作單：Rubin Ultra 比例讀 Siliconomics Units；以每 GW GPU 數換成 GW 占比（GL01-11）；2027 Blackwell 2.4M 顆不計（工作單：2027 起 VR200）")
    r += 1
    NNVR = r
    PL.cell(r, 1, "市場新增 GW 非 NVIDIA 占比"); PL.cell(r, 2, "nnv_mkt"); PL.cell(r, 3, "%")
    ir = INROW["nnv_mkt"]
    for y in YEARS:
        inref = f"In!${IC[y]}${ir}"
        if 2026 <= y <= 2029:
            amd = "SIL_it_AMD8" if y == 2026 else "SIL_it_AMDR"
            nnv = f"({sv('SIF_UnitsPlat_amd', y)}*{amd}+{sv('SIF_UnitsPlat_tpu', y)}*SIL_it_TPU+{sv('SIF_UnitsPlat_trn', y)}*SIL_it_TRN+{sv('SIF_UnitsPlat_oth', y)}*SIL_it_ASIC)"
            nv = f"({sv('SIF_UnitsPlat_bw', y)}*SIL_it_GB+{sv('SIF_UnitsPlat_rb', y)}*SIL_it_VR+{sv('SIF_UnitsPlat_ru', y)}*SIL_it_RU+{sv('SIF_UnitsPlat_hp', y)}*SIL_it_HP)"
            d = f"{nnv}/({nnv}+{nv})"
        elif y < 2026:
            d = f"{YC[2026]}{r}"
        else:
            d = f"{YC[2029]}{r}"
        PL.cell(r, YEARS.index(y) + 4, f'=IF({inref}<>"",{inref},{d})').number_format = "0.0%"
    PL.cell(r, 14, "Derived：Siliconomics Units 非 NVIDIA 平台 IT capex 占比（工作單第 2 節）；2023–25 沿用 2026（待決）、2030–32 沿用 2029")
    r += 1
    FACR = r
    PL.cell(r, 1, "每 IT GW 對應設施電力（TK IF_FacilityGW，世代加權）"); PL.cell(r, 3, "x")
    for y in YEARS:
        c = YC[y]
        PL.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"{c}{MIXR[g]}*TK_FacilityGW_{g}" for g in GENS)).number_format = "0.000"
    PL.cell(r, 14, "TK_In（Tokenomics v21.7 Interface 第 8 列；快照未含，直接讀活頁簿，SRC_GL_006）；非 NVIDIA 以 NVIDIA 世代加權近似")
    r += 1
    MULT = {}
    for p in PK:
        if p not in CHIP: continue
        cr = CHIP[p]; c25 = f"In!${IC[2025]}${cr}"
        ov = seed["ovr"].get(p)
        PL.cell(r, 1, f"晶片 → IT 倍數：{PN[p]}"); PL.cell(r, 2, f"mult_{p}"); PL.cell(r, 3, "x")
        if ov:
            f0, flo, fhi = f"={ov}/{c25}", f"={ov}_lo/{c25}", f"={ov}_hi/{c25}"
            PL.cell(r, 14, f"C 段 2025 更正值 {ov} ÷ 2025 晶片功耗（D＝基準、E＝低、F＝高）")
        else:
            f0, flo, fhi = "=k_chip2it", "=k_chip2it_lo", "=k_chip2it_hi"
            PL.cell(r, 14, "k_chip2it（D＝基準、E＝低、F＝高）")
        for j, f in ((4, f0), (5, flo), (6, fhi)):
            PL.cell(r, j, f).number_format = "0.000"
        MULT[p] = (f"Players!$D${r}", f"Players!$E${r}", f"Players!$F${r}"); r += 1
    r += 1
    PR = {}
    OBS = {p for p in PK if f"obs_{p}" in INROW or f"obschip_{p}" in INROW}
    CAMPSUM = lambda p, y: f'SUMIFS(Pipeline!$L${{c0}}:$L${{c1}},Pipeline!$C${{c0}}:$C${{c1}},"{p}",Pipeline!$Q${{c0}}:$Q${{c1}},1,Pipeline!$P${{c0}}:$P${{c1}},{y})'
    PLAYER_CELLS = []   # (row, col, formula template needing c0/c1)
    RNG = {"neo": {2025: "rng25_neo"}, "cn": {2025: "rng25_cn", 2026: "rng26_cn"}}
    for p in PK:
        ir = INROW[f"mw_{p}"]
        PL.cell(r, 1, PN[p]).font = BOLD; PL.cell(r, 2, p); PL.cell(r, 14, f"In 第 {ir} 列；口徑＝In!C{ir}")
        r += 1
        rows = {}
        for k, lab, unit in (("in", "輸入（IT GW；設施 ÷ IF_FacilityGW）", "GW"), ("chip", "晶片換算 IT GW", "GW"), ("k", "已知值", "GW"),
                             ("L", "前一已知值", "GW"), ("LY", "前一已知年", "年"), ("N", "下一已知值", "GW"), ("NY", "下一已知年", "年"),
                             ("cp", "當年 T2 園區上線（D0，IT GW）", "GW"), ("tk", "年底上線 IT GW", "GW"), ("f", "旗標", "—"),
                             ("lo", "年底上線 IT GW：低", "GW"), ("hi", "年底上線 IT GW：高", "GW"),
                             ("alt", "另讀（口徑未說明者讀為設施 ÷ IF_FacilityGW）", "GW"), ("avg", "年平均 IT GW", "GW"),
                             ("add", "年新增 IT GW", "GW"), ("stk", "新納入存量 IT GW", "GW"), ("nnv", "新增 GW 非 NVIDIA 占比", "%")):
            rows[k] = r; PL.cell(r, 1, "　" + lab); PL.cell(r, 2, f"{p}_{k}"); PL.cell(r, 3, unit); r += 1
        for i, y in enumerate(YEARS):
            c = YC[y]; inref = f"In!${IC[y]}${ir}"; col = i + 4
            pc = YC.get(y - 1); nc = YC.get(y + 1)
            if p == "oth" and y == 2025 and "other_incl_china" in INROW:
                PL.cell(rows["in"], col, f'=IF({inref}<>"",{inref}/1000,IF(other_incl_china=1,oth_excn,oth_full)/1000)').number_format = "0.000"
            else:
                PL.cell(rows["in"], col, f'=IF({inref}="","",IF(In!$C${ir}="設施",{inref}/1000/{c}${FACR},{inref}/1000))').number_format = "0.000"
            if p in CHIP:
                cref = f"In!${IC[y]}${CHIP[p]}"
                PL.cell(rows["chip"], col, f'=IF({cref}="","",{cref}*{MULT[p][0]}/1000)').number_format = "0.000"
            else:
                PL.cell(rows["chip"], col, '=""')
            PLAYER_CELLS.append((rows["cp"], col, "=" + CAMPSUM(p, y)))
            IN_, CH_ = f"{c}{rows['in']}", f"{c}{rows['chip']}"
            if p in OBS:
                if f"obschip_{p}" in INROW:
                    ov = f"obschip_{p}*{MULT[p][0]}/1000"
                else:
                    ov = f"obs_{p}/1000"
                after = (f'SUMIFS(Pipeline!$L${{c0}}:$L${{c1}},Pipeline!$C${{c0}}:$C${{c1}},"{p}",Pipeline!$Q${{c0}}:$Q${{c1}},1,'
                         f'Pipeline!$K${{c0}}:$K${{c1}},">"&obsd_{p},Pipeline!$K${{c0}}:$K${{c1}},"<="&{y + 1})')
                kf = f'=IF({IN_}<>"",{IN_},IF({CH_}<>"",{CH_},IF(INT(obsd_{p})={y},{ov}+{after},"")))'
                PLAYER_CELLS.append((rows["k"], col, kf))
            else:
                PL.cell(rows["k"], col, f'=IF({IN_}<>"",{IN_},IF({CH_}<>"",{CH_},""))')
            K = f"{c}{rows['k']}"
            if pc is None:
                PL.cell(rows["L"], col, f'=IF({K}<>"",{K},"")'); PL.cell(rows["LY"], col, f'=IF({K}<>"",{y},"")')
            else:
                PL.cell(rows["L"], col, f'=IF({K}<>"",{K},{pc}{rows["L"]})'); PL.cell(rows["LY"], col, f'=IF({K}<>"",{y},{pc}{rows["LY"]})')
            if nc is None:
                PL.cell(rows["N"], col, f'=IF({K}<>"",{K},"")'); PL.cell(rows["NY"], col, f'=IF({K}<>"",{y},"")')
            else:
                PL.cell(rows["N"], col, f'=IF({K}<>"",{K},{nc}{rows["N"]})'); PL.cell(rows["NY"], col, f'=IF({K}<>"",{y},{nc}{rows["NY"]})')
            Lr, LYr, Nr, NYr, CP_ = (f"{c}{rows[x]}" for x in ("L", "LY", "N", "NY", "cp"))
            prevv = f"{pc}{rows['tk']}" if pc else '""'
            ext = f'IF({prevv}<>"",{prevv}+{CP_},"")' if pc else '""'
            if p == "cn" and 2027 <= y <= 2030:
                ext = (f'IF({YC[2026]}{rows["tk"]}="","",{YC[2026]}{rows["tk"]}*(cn_tot26*(cn_tot30/cn_tot26)^(({y}-2026)/4))*(cn_sh26+(cn_sh30-cn_sh26)*({y}-2026)/4)'
                       f'/(cn_tot26*cn_sh26))')
            elif p == "cn" and y > 2030:
                ext = f'IF({prevv}<>"",{prevv}*(1+g_lt),"")'
            cc = PL.cell(rows["tk"], col, f'=IF({K}<>"",{K},IF(AND({Lr}<>"",{Nr}<>""),{Lr}+({Nr}-{Lr})*({y}-{LYr})/({NYr}-{LYr}),{ext}))')
            cc.number_format = "0.000"; cc.fill = KEY
            PL.cell(rows["f"], col, f'=IF({IN_}<>"",1,IF({CH_}<>"",5,IF({K}<>"",2,IF(AND({Lr}<>"",{Nr}<>""),3,IF({c}{rows["tk"]}<>"",4,0)))))')
            T = f"{c}{rows['tk']}"
            # 低／高
            lo, hi = T, T
            if p in CHIP:
                cref = f"In!${IC[y]}${CHIP[p]}"
                lo = f'IF(AND({c}{rows["f"]}=5,{cref}<>""),{cref}*{MULT[p][1]}/1000,{T})'
                hi = f'IF(AND({c}{rows["f"]}=5,{cref}<>""),{cref}*{MULT[p][2]}/1000,{T})'
                if p == "msft" and y == 2025:
                    lo = f"MIN({lo},msft_guardian_lo/1000)"
            if p in OBS and any(sc.get("key") == f"obs_{p}" and sc.get("lo") is not None for sc in seed["scalars"]):
                lo = f'IF({c}{rows["f"]}=2,{T}-(obs_{p}-obs_{p}_lo)/1000,{lo})'
                hi = f'IF({c}{rows["f"]}=2,{T}+(obs_{p}_hi-obs_{p})/1000,{hi})'
            if p == "oth" and y == 2025 and "other_incl_china" in INROW:
                lo = "IF(other_incl_china=1,oth_excn_lo,oth_full_lo)/1000"; hi = "IF(other_incl_china=1,oth_excn_hi,oth_full_hi)/1000"
            elif p in RNG and y in RNG[p]:
                nm_ = RNG[p][y]; div = f"/{c}${FACR}" if p in ("cn", "oth") else ""
                lo = f"{nm_}_lo/1000{div}"; hi = f"{nm_}_hi/1000{div}"
            PL.cell(rows["lo"], col, f'=IF({T}="","",{lo})').number_format = "0.000"
            PL.cell(rows["hi"], col, f'=IF({T}="","",{hi})').number_format = "0.000"
            PL.cell(rows["alt"], col, f'=IF({T}="","",IF(In!$C${ir}="未說明",{T}/{c}${FACR},{T}))').number_format = "0.000"
            if pc is None:
                PL.cell(rows["avg"], col, '=""'); PL.cell(rows["add"], col, 0)
                PL.cell(rows["stk"], col, f'=IF({T}<>"",{T},0)').number_format = "0.000"
            else:
                PT = f"{pc}{rows['tk']}"
                PL.cell(rows["avg"], col, f'=IF(AND({PT}<>"",{T}<>""),({PT}+{T})/2,"")').number_format = "0.000"
                PL.cell(rows["add"], col, f'=IF(AND({PT}<>"",{T}<>""),{T}-{PT},0)').number_format = "0.000"
                PL.cell(rows["stk"], col, f'=IF(AND({PT}="",{T}<>""),{T},0)').number_format = "0.000"
            PL.cell(rows["nnv"], col, (f"=nnv_{p}" if f"nnv_{p}" in INROW else f"={c}${NNVR}")).number_format = "0.0%"
        PR[p] = rows; r += 1
    PL.cell(r, 1, "市場合計（Σ 營運者；主要玩家＋其他）").font = BOLD; r += 1
    TOT = {}
    for k, lab, key in (("tk", "年底上線 IT GW（＝TK GW）", "ye"), ("lo", "年底上線 IT GW：低", "lo"), ("hi", "年底上線 IT GW：高", "hi"),
                        ("alt", "年底上線 IT GW（另讀：口徑未說明者讀為設施）", "alt"), ("fac", "年底上線 設施 GW（IT × IF_FacilityGW）", "fac"),
                        ("avg", "年平均 IT GW", "avg"), ("add", "年新增 IT GW", "add"), ("stk", "新納入存量 IT GW", "stk"),
                        ("cnx", "年底上線 IT GW（不含中國）", "cnx"), ("cnt", "有值玩家數", "cnt")):
        PL.cell(r, 1, lab); PL.cell(r, 2, f"mkt_{key}")
        for i, y in enumerate(YEARS):
            c = YC[y]
            if key == "fac":
                f = f"={c}{TOT['ye']}*{c}${FACR}"
            elif key == "avg":
                f = '=""' if i == 0 else f"=({YC[y - 1]}{TOT['ye']}+{c}{TOT['ye']})/2"
            elif key == "cnt":
                f = "=COUNT(" + ",".join(f"{c}{PR[p]['tk']}" for p in PK) + ")"
            elif key == "cnx":
                f = f"={c}{TOT['ye']}-N({c}{PR['cn']['tk']})"
            else:
                f = "=SUM(" + ",".join(f"{c}{PR[p][k]}" for p in PK) + ")"
            cc = PL.cell(r, i + 4, f); cc.number_format = "0" if key == "cnt" else "0.000"
            if key == "ye": cc.fill = KEY
        TOT[key] = r; r += 1
    PL_LAST = r
    PL.column_dimensions["A"].width = 52; PL.column_dimensions["B"].width = 12; PL.column_dimensions["N"].width = 60
    for y in YEARS: PL.column_dimensions[YC[y]].width = 9

    # ================= Matrix：營運者 × 使用者 =================
    MX["A1"] = ("Matrix — 營運者 × 使用者（年底 IT GW）。列＝營運者（Players 年底值），欄＝使用者：In T1 每筆分配（比＝× 營運者年底 GW；GW＝絕對值）加總；"
                "『未分配』＝列和缺口（無揭露者；Decisions GL01-15）。列和＝Players；欄和合計＝總量（市場總量只按營運者加總一次）")
    r = 3; MXR = {}
    NU = len(U); UNC = 3 + NU; SUMC = UNC + 1; PLC = UNC + 2; DFC = UNC + 3
    L = get_column_letter
    for yi, y in enumerate((2025, 2026, 2027)):
        vcol = L(11 + yi)   # In T1：K／L／M
        MX.cell(r, 1, f"{y} 年底").font = BOLD; r += 1
        hdr(MX, r, ["營運者", "鍵"] + [u["name"] for u in U] + ["未分配", "列和", "Players", "差"]); r += 1
        r0 = r
        for p in PK:
            MX.cell(r, 1, PN[p]); MX.cell(r, 2, p)
            MX.cell(r, PLC, f"=N(Players!{YC[y]}{PR[p]['tk']})").font = GREEN
            P_ = f"${L(PLC)}{r}"
            for j, u in enumerate(U):
                f = (f'=SUMPRODUCT((In!$H${AL0}:$H${AL1}="{p}")*(In!$I${AL0}:$I${AL1}="{u["key"]}")*'
                     f'((In!$J${AL0}:$J${AL1}="比")*{P_}+(In!$J${AL0}:$J${AL1}<>"比"))*In!{vcol}${AL0}:{vcol}${AL1})')
                MX.cell(r, 3 + j, f).number_format = "0.000"
            MX.cell(r, UNC, f"={P_}-SUM(C{r}:{L(UNC - 1)}{r})").number_format = "0.000"
            MX.cell(r, SUMC, f"=SUM(C{r}:{L(UNC)}{r})").number_format = "0.000"
            MX.cell(r, DFC, f"={L(SUMC)}{r}-{L(PLC)}{r}").number_format = "0.000000"
            MX.cell(r, PLC).number_format = "0.000"
            r += 1
        r1 = r - 1
        MX.cell(r, 1, "欄和（使用者合計）").font = BOLD
        for j in range(3, DFC + 1):
            col = L(j)
            MX.cell(r, j, f"=SUM({col}{r0}:{col}{r1})").number_format = "0.000"
        MXR[y] = dict(r0=r0, r1=r1, sum=r); r += 1
        for ukey, uk in (("usr_oai", "oai"), ("usr_ant", "ant")):
            ir = INROW[ukey]; ucol = 3 + [u["key"] for u in U].index(uk); UC = L(ucol)
            MX.cell(r, 1, f"對照：{ukey} 報導值（口徑 In!C{ir}；未說明讀為 IT）")
            MX.cell(r, ucol, f'=IF(In!{IC[y]}{ir}="","",IF(In!$C${ir}="設施",In!{IC[y]}{ir}/Players!{YC[y]}${FACR},In!{IC[y]}{ir}))').number_format = "0.000"
            MX.cell(r, SUMC, f'=IF(OR({UC}{r}="",N({UC}{r})=0),"",{UC}{MXR[y]["sum"]}/{UC}{r})').number_format = "0.00"
            MX.cell(r, PLC, "← 矩陣欄和 ÷ 報導（< 1＝揭露不足，差額在『未分配』）")
            r += 1
        r += 1
    MX_LAST = r
    MX.column_dimensions["A"].width = 46
    for j in range(3, DFC + 1): MX.column_dimensions[L(j)].width = 10

    # ================= Pipeline：園區管線 × 延遲情景 =================
    PP["A1"] = ("Pipeline — 上線事件 × 延遲情景（IT GW）。事件＝(1) 園區（In T2，有上線年月且納入＝1）＋(2) 各玩家各年『未分園區殘差』＝Players 年新增 − 該年 T2 園區；"
                "＋新納入存量（不延遲）。時間為小數年（年 Y 月 M 底＝Y＋M/12；Y 年底＝Y＋1）。延遲只作用於評價時點 asof 之後：D0＝0；"
                "D1＝每筆排定量的 dly_d1_s 延 dly_d1_m 月（園區有已知延遲者全額延已知月數）；D2＝D1 之上，原排定年在 dly_d2_y0–dly_d2_y1 且未確保電力者，"
                "dly_d2_bind 再延 dly_d2_m 月（殘差視為未確保電力）。殘差在年內均勻上線（觀測年在觀測時點前已上線）")
    r = 3
    PP.cell(r, 1, "A. 園區（In T2 鏡像與上線時點）").font = BOLD; r += 1
    hdr(PP, r, ["園區鍵", "名稱", "營運者", "使用者", "MW", "口徑", "納入", "上線年", "上線月", "已知延遲月", "t 公告", "IT GW", "t D0", "t D1 未延", "t D1 延後",
                "D0 上線年", "有效納入", "D2 適用", "已確保電力"]); r += 1
    CR0 = r
    for k, ir in CAMP:
        PP.cell(r, 1, k); PP.cell(r, 2, f"=In!B{ir}").font = GREEN
        for j, col in ((3, "H"), (4, "I"), (5, "J"), (7, "N"), (8, "K"), (9, "L"), (19, "O")):
            PP.cell(r, j, f'=IF(In!{col}{ir}="","",In!{col}{ir})').font = GREEN
        PP.cell(r, 10, f"=N(In!M{ir})").font = GREEN   # 已知延遲月：一律數值（空白＝0；文字 "" 在比較中會被當成大於 0）
        PP.cell(r, 6, f"=In!C{ir}").font = GREEN
        PP.cell(r, 11, f'=IF(OR(H{r}="",I{r}=""),9999,H{r}+I{r}/12)').number_format = "0.000"
        PP.cell(r, 12, f'=IF(E{r}="",0,IF(F{r}="設施",E{r}/1000/INDEX(Players!$D${FACR}:$M${FACR},1,MAX(1,MIN(10,H{r}-2022))),E{r}/1000*g_ITtoFac))').number_format = "0.000"
        PP.cell(r, 13, f"=K{r}").number_format = "0.000"
        PP.cell(r, 14, f"=IF(K{r}>asof,K{r}+N(J{r})/12,K{r})").number_format = "0.000"
        PP.cell(r, 15, f"=IF(K{r}>asof,K{r}+IF(N(J{r})>0,N(J{r}),dly_d1_m)/12,K{r})").number_format = "0.000"
        PP.cell(r, 16, f'=IF(K{r}=9999,"",ROUNDUP(K{r},0)-1)')
        PP.cell(r, 17, f"=IF(AND(N(G{r})=1,K{r}<9999,L{r}>0),1,0)")
        PP.cell(r, 18, f'=IF(AND(K{r}>asof,K{r}<9999,N(S{r})<>1,P{r}>=dly_d2_y0,P{r}<=dly_d2_y1),1,0)')
        r += 1
    CR1 = r - 1
    for (rr_, cc_, f) in PLAYER_CELLS:
        PL.cell(rr_, cc_, f.replace("{c0}", str(CR0)).replace("{c1}", str(CR1))).number_format = "0.000"
    rng = lambda col: f"${col}${CR0}:${col}${CR1}"
    r += 1
    PP.cell(r, 1, "B. 未分園區殘差（IT GW；Players 年新增 − 該營運者該年 D0 上線的 T2 園區）").font = BOLD; r += 1
    hdr(PP, r, ["玩家", "鍵", "類別"] + [str(y) for y in YEARS]); r += 1
    RES = {}; ROBS = {}; RUNI = {}
    for kind, store in (("殘差合計", RES), ("觀測型（不延遲）", ROBS), ("均勻型（延遲適用）", RUNI)):
        for p in PK:
            PP.cell(r, 1, PN[p]); PP.cell(r, 2, p); PP.cell(r, 3, kind)
            for y in OUTY:
                c = YC[y]
                if store is RES:
                    f = f'=Players!{c}{PR[p]["add"]}-SUMIFS({rng("L")},{rng("C")},"{p}",{rng("Q")},1,{rng("P")},{y})'
                elif store is ROBS:
                    f = f"=IF(Players!{c}{PR[p]['f']}=2,{c}{RES[p]},0)"
                else:
                    f = f"={c}{RES[p]}-{c}{ROBS[p]}"
                PP.cell(r, YEARS.index(y) + 4, f).number_format = "0.000"
            store[p] = r; r += 1
        r += 1
    AGG = {}
    for key, lab, f_ in (("res", "殘差合計", lambda c: f"=SUM({c}{RES[PK[0]]}:{c}{RES[PK[-1]]})"),
                         ("robs", "觀測型殘差合計", lambda c: f"=SUM({c}{ROBS[PK[0]]}:{c}{ROBS[PK[-1]]})"),
                         ("runi", "均勻型殘差合計", lambda c: f"=SUM({c}{RUNI[PK[0]]}:{c}{RUNI[PK[-1]]})"),
                         ("camp", "園區（D0 上線年）合計", None),
                         ("stk", "新納入存量（Players）", lambda c: f"=Players!{c}{TOT['stk']}"),
                         ("addchk", "殘差＋園區（須＝Players 年新增）", None),
                         ("pwr", "電力證據路徑：T2 納入＝2 的供電合約列 D0 上線（IT GW；不入 Players，只用於 WIF_Coverage_Power）", None)):
        PP.cell(r, 1, lab).font = BOLD
        for y in YEARS:
            c = YC[y]
            if key == "camp":
                f = f'=SUMIFS({rng("L")},{rng("Q")},1,{rng("P")},{y})'
            elif key == "addchk":
                f = f"={c}{AGG['res']}+{c}{AGG['camp']}"
            elif key == "pwr":
                f = f'=SUMIFS({rng("L")},{rng("G")},2,{rng("P")},{y})'
            elif y == 2023 and key in ("res", "robs", "runi"):
                f = 0
            else:
                f = f_(c)
            PP.cell(r, YEARS.index(y) + 4, f).number_format = "0.000"
        AGG[key] = r; r += 1
    r += 1
    PP.cell(r, 1, "C. 均勻型事件的年底上線比例矩陣（列＝年底 z，欄＝事件年 y；asof 前部分不延遲，asof 後部分依情景延後）").font = BOLD; r += 1
    FRAC = {}
    for D in DLY:
        hdr(PP, r, [f"延遲情景 {D}", "年底 z", ""] + [""] + [str(y) for y in OUTY]); r += 1
        FRAC[D] = {}
        for z in YEARS:
            PP.cell(r, 1, f"{D}：{z} 年底"); PP.cell(r, 2, z)
            for y in OUTY:
                p1 = f"MAX(0,MIN({z}+1,{y}+MAX(0,MIN(1,asof-{y})))-{y})"
                def p2(d):
                    return f"MAX(0,MIN({z}+1,{y}+1+{d})-(MAX({y},asof)+{d}))"
                d1 = f"((1-dly_d1_s)*{p2('0')}+dly_d1_s*{p2('dly_d1_m/12')})"
                d1s = f"((1-dly_d1_s)*{p2('dly_d2_m/12')}+dly_d1_s*{p2('(dly_d1_m+dly_d2_m)/12')})"
                e = f"AND({y}>=dly_d2_y0,{y}<=dly_d2_y1)"
                if D == "D0":
                    f = f"={p1}+{p2('0')}"
                elif D == "D1":
                    f = f"={p1}+{d1}"
                else:
                    f = f"={p1}+IF({e},(1-dly_d2_bind)*{d1}+dly_d2_bind*{d1s},{d1})"
                PP.cell(r, YEARS.index(y) + 4, f).number_format = "0.00"
            FRAC[D][z] = r; r += 1
        r += 1
    PP.cell(r, 1, "D. 年底上線 IT GW（延遲情景）").font = BOLD; r += 1
    hdr(PP, r, ["項目", "情景", ""] + [str(y) for y in YEARS]); r += 1
    PA = {}
    def prow(key, lab, D, fz):
        nonlocal r
        PP.cell(r, 1, lab); PP.cell(r, 2, D)
        for z in YEARS:
            cc = PP.cell(r, YEARS.index(z) + 4, fz(z)); cc.number_format = "0.000"
        PA[key] = r; r += 1
    prow("stkc", "新納入存量累計", "", lambda z: f"=SUM($D${AGG['stk']}:{YC[z]}{AGG['stk']})")
    prow("obsc", "觀測型殘差累計（不延遲）", "", lambda z: f"=SUM($D${AGG['robs']}:{YC[z]}{AGG['robs']})")
    for D in DLY:
        prow("uni" + D, "均勻型殘差（延遲後）", D, lambda z, D=D: f"=SUMPRODUCT($E${AGG['runi']}:$M${AGG['runi']},$E${FRAC[D][z]}:$M${FRAC[D][z]})")
    L_, Q_, M_, N_, O_, R_, J_ = (rng(x) for x in "LQMNORJ")
    def c1(z, sh=""):
        # D1 園區貢獻：已知延遲者全額在 O（＝N）；否則 (1−s) 在 N、s 在 O
        return (f"((({J_}>0)+({J_}<=0)*(1-dly_d1_s))*({N_}{sh}<={z}+1)+({J_}<=0)*dly_d1_s*({O_}{sh}<={z}+1))")
    prow("campD0", "園區（延遲後）", "D0", lambda z: f"=SUMPRODUCT({L_},{Q_},({M_}<={z}+1)*1)")
    prow("campD1", "園區（延遲後）", "D1", lambda z: f"=SUMPRODUCT({L_},{Q_},{c1(z)})")
    prow("campD2", "園區（延遲後）", "D2", lambda z: f"=SUMPRODUCT({L_},{Q_},(1-dly_d2_bind*{R_})*{c1(z)}+dly_d2_bind*{R_}*{c1(z, '+dly_d2_m/12')})")
    for D in DLY:
        prow("ye" + D, "年底上線 IT GW", D, lambda z, D=D: f"={YC[z]}{PA['stkc']}+{YC[z]}{PA['obsc']}+{YC[z]}{PA['uni' + D]}+{YC[z]}{PA['camp' + D]}")
        PP.cell(PA["ye" + D], 1).font = BOLD
        for z in YEARS: PP.cell(PA["ye" + D], YEARS.index(z) + 4).fill = KEY
    for D in DLY:
        prow("add" + D, "年新增 IT GW（淨，不含新納入存量）", D,
             lambda z, D=D: "=0" if z == 2023 else f"={YC[z]}{PA['ye' + D]}-{YC[z - 1]}{PA['ye' + D]}-{YC[z]}{AGG['stk']}")
    for D in ("D1", "D2"):
        prow("sno" + D, "已出貨未上線 IT GW（相對 D0）", D, lambda z, D=D: f"={YC[z]}{PA['yeD0']}-{YC[z]}{PA['ye' + D]}")
    PP_LAST = r
    PP.column_dimensions["A"].width = 40; PP.column_dimensions["B"].width = 26
    for col in "CDEFGHIJKLMNOPQRS": PP.column_dimensions[col].width = 9

    # ================= CapexGW：甲2 資本支出推得 GW =================
    CX["A1"] = ("CapexGW — 甲2：資本支出推得 GW（MSFT、GOOGL、AMZN、META、ORCL、CRWV、NBIS）。財年資本支出（In Y2；個股模型 annualEstimates／指引，MAG 讀現行檔：microsoft 2026-10-11、"
                "alphabet 2026-10-09、amazon 2026-10-08 版，PR #190 未合併）→ 補值（最早已知年之前不回填，留空並跳過該年 甲2；最後年後 ×（1＋cap_g_ext））→ 曆年（財年結束月拆分）→ × AI 占比 → "
                "當年 GW 資金＝IT 占比 × 當年＋廠房占比 ×（w0 當年＋w1 前一年＋w2 前兩年；lag_bldg）→ ÷ 每 IT GW 合計資本支出（TK IF_CapexTotal 世代加權；NonNV × 持有比 × VR200）＝年新增 IT GW。"
                "[Derived]；輸入為 Interested-party／共識")
    CX["A2"] = "主路徑：2024–2026 用 甲（Players：園區＋實測）；2027 起本頁七家改用 甲2，其餘玩家仍用 甲。Decisions GL01-19、GL01-21（廠房多由房東出資者可能低估）"
    hdr(CX, 3, ["項目", "玩家", "單位"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    CX.cell(r, 1, "廠房提前入帳權重 w0／w1／w2（lag_bldg）")
    CX.cell(r, 4, "=MAX(0,1-lag_bldg/12)"); CX.cell(r, 5, "=MAX(0,MIN(lag_bldg/12,2-lag_bldg/12))"); CX.cell(r, 6, "=MAX(0,MIN(1,lag_bldg/12-1))")
    XW = (f"CapexGW!$D${r}", f"CapexGW!$E${r}", f"CapexGW!$F${r}"); r += 1
    XWV = {}
    for sx, nm_ in (("lo", "lag_bldg_lo"), ("hi", "lag_bldg_hi")):
        CX.cell(r, 1, f"敏感度：w0／w1／w2（{nm_}）")
        CX.cell(r, 4, f"=MAX(0,1-{nm_}/12)"); CX.cell(r, 5, f"=MAX(0,MIN({nm_}/12,2-{nm_}/12))"); CX.cell(r, 6, f"=MAX(0,MIN(1,{nm_}/12-1))")
        XWV[sx] = (f"CapexGW!$D${r}", f"CapexGW!$E${r}", f"CapexGW!$F${r}"); r += 1
    CX.cell(r, 1, "NVIDIA 每 IT GW IT capex（世代加權）"); CX.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; CX.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexIT_{g}" for g in GENS)).number_format = "0.00"
    XIT = r; r += 1
    CX.cell(r, 1, "NVIDIA 每 IT GW 廠房 capex（世代加權）"); CX.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; CX.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexFacility_{g}" for g in GENS)).number_format = "0.00"
    XFAC = r; r += 2
    PLAT = {p_["key"]: p_["nnv_plat"] for p_ in P}
    CAPP = ["msft", "googl", "amzn", "meta", "orcl", "crwv", "nbis"]
    XR = {}
    for p in CAPP:
        cr = INROW[f"cap_{p}"]; fye = FYE.get(p, "12"); aik = f"ai_{p}" if f"ai_{p}" in INROW else "ai_sh"
        CX.cell(r, 1, PN[p]).font = BOLD; CX.cell(r, 2, p); CX.cell(r, 14, f"In 第 {cr} 列（財年）；AI 占比＝{aik}；財年結束月＝{fye}")
        r += 1
        rows = {}
        for k, lab, unit in (("in", "財年資本支出（In Y2）", "US$B"), ("cnt", "　已知值累計個數", "—"), ("N", "　下一已知值", "US$B"), ("fy", "財年資本支出（補值）", "US$B"),
                             ("cy", "曆年資本支出", "US$B"), ("ai", "曆年 AI 資本支出", "US$B"), ("pit", "每 IT GW IT capex", "US$B"), ("pfac", "每 IT GW 廠房 capex", "US$B"),
                             ("fund", "當年上線 GW 的資金（廠房提前入帳）", "US$B"), ("gw", "甲2 年新增 IT GW", "GW"), ("site", "甲 年新增 IT GW（Players）", "GW"),
                             ("main", "主路徑 年新增 IT GW（2027 起＝甲2）", "GW"), ("ye", "主路徑 年底 IT GW", "GW"),
                             ("s_ailo", "敏感度：甲2 年新增（AI 占比 低）", "GW"), ("s_aihi", "敏感度：甲2 年新增（AI 占比 高）", "GW"),
                             ("s_laglo", "敏感度：甲2 年新增（lag_bldg 低）", "GW"), ("s_laghi", "敏感度：甲2 年新增（lag_bldg 高）", "GW")):
            rows[k] = r; CX.cell(r, 1, "　" + lab); CX.cell(r, 2, f"{p}_{k}"); CX.cell(r, 3, unit); r += 1
        for i, y in enumerate(YEARS):
            c = YC[y]; col = i + 4; pc = YC.get(y - 1); nc = YC.get(y + 1)
            inref = f"In!{IC[y]}{cr}"
            CX.cell(rows["in"], col, f'=IF({inref}="","",{inref})').font = GREEN
            IN_ = f"{c}{rows['in']}"
            CX.cell(rows["cnt"], col, f'=({IN_}<>"")' + (f"+{pc}{rows['cnt']}" if pc else "")).number_format = "0"
            CX.cell(rows["N"], col, f'=IF({IN_}<>"",{IN_},' + (f"{nc}{rows['N']})" if nc else '"")'))
            prevfy = f"{pc}{rows['fy']}" if pc else '""'
            prevcnt = f"{pc}{rows['cnt']}" if pc else "0"
            CX.cell(rows["fy"], col, f'=IF({IN_}<>"",{IN_},IF({prevcnt}>0,{prevfy}*(1+cap_g_ext),""))').number_format = "0.0"
            FY_ = f"{c}{rows['fy']}"
            nfy = f"{nc}{rows['fy']}" if nc else f'IF({FY_}="","",{FY_}*(1+cap_g_ext))'
            CX.cell(rows["cy"], col, f'=IF(OR({FY_}="",AND({fye}<12,{nfy}="")),"",{FY_}*{fye}/12+IF({fye}<12,{nfy},0)*(12-{fye})/12)').number_format = "0.0"
            CX.cell(rows["ai"], col, f'=IF({c}{rows["cy"]}="","",{c}{rows["cy"]}*{aik})').number_format = "0.0"
            nnv = f"Players!{c}{PR[p]['nnv']}"
            CX.cell(rows["pit"], col, f"=(1-{nnv})*{c}${XIT}+{nnv}*TK_NonNV_Hold_{PLAT[p]}*TK_CapexIT_VR200").number_format = "0.00"
            CX.cell(rows["pfac"], col, f"=(1-{nnv})*{c}${XFAC}+{nnv}*TK_NonNV_Hold_{PLAT[p]}*TK_CapexFacility_VR200").number_format = "0.00"
            PIT, PFA = f"{c}{rows['pit']}", f"{c}{rows['pfac']}"
            terms = [f"{XW[0]}*{c}{rows['ai']}"]
            if YC.get(y - 1): terms.append(f"{XW[1]}*{YC[y - 1]}{rows['ai']}")
            else: terms.append(f"{XW[1]}*{c}{rows['ai']}")
            if YC.get(y - 2): terms.append(f"{XW[2]}*{YC[y - 2]}{rows['ai']}")
            else: terms.append(f"{XW[2]}*{c}{rows['ai']}")
            need = [f"{c}{rows['ai']}", f"{(YC.get(y - 1) or c)}{rows['ai']}", f"{(YC.get(y - 2) or c)}{rows['ai']}"]
            miss = "OR(" + ",".join(f'{x}=""' for x in need) + ")"
            CX.cell(rows["fund"], col, f'=IF({miss},"",{PIT}/({PIT}+{PFA})*{c}{rows["ai"]}+{PFA}/({PIT}+{PFA})*(' + "+".join(terms) + "))").number_format = "0.0"
            cc = CX.cell(rows["gw"], col, f'=IF({c}{rows["fund"]}="","",{c}{rows["fund"]}/({PIT}+{PFA}))'); cc.number_format = "0.000"; cc.fill = KEY
            G_ = f"{c}{rows['gw']}"
            CX.cell(rows["s_ailo"], col, f'=IF({G_}="","",{G_}*{aik}_lo/{aik})').number_format = "0.000"
            CX.cell(rows["s_aihi"], col, f'=IF({G_}="","",{G_}*{aik}_hi/{aik})').number_format = "0.000"
            for sx in ("lo", "hi"):
                W = XWV[sx]
                tt = [f"{W[0]}*{need[0]}", f"{W[1]}*{need[1]}", f"{W[2]}*{need[2]}"]
                CX.cell(rows["s_lag" + sx], col, f'=IF({miss},"",({PIT}/({PIT}+{PFA})*{c}{rows["ai"]}+{PFA}/({PIT}+{PFA})*(' + "+".join(tt) + f"))/({PIT}+{PFA}))").number_format = "0.000"
            CX.cell(rows["site"], col, f"=Players!{c}{PR[p]['add']}").font = GREEN
            CX.cell(rows["main"], col, f'=IF(AND({y}>=2027,{G_}<>""),{G_},{c}{rows["site"]})').number_format = "0.000"
            if pc is None:
                CX.cell(rows["ye"], col, f"=N(Players!{c}{PR[p]['tk']})").number_format = "0.000"
            else:
                CX.cell(rows["ye"], col, f"=IF({y}>=2027,{pc}{rows['ye']}+{c}{rows['main']},N(Players!{c}{PR[p]['tk']}))").number_format = "0.000"
        XR[p] = rows; r += 1
    CX.cell(r, 1, "七家合計").font = BOLD; r += 1
    XT = {}
    for k, lab in (("gw", "甲2 年新增 IT GW（缺值以 甲 代入）"), ("site", "甲 年新增 IT GW"), ("adj", "主路徑調整（2027–2029：甲2 − 甲）"),
                   ("s_ailo", "敏感度：甲2 年新增（AI 占比 低）"), ("s_aihi", "敏感度：甲2 年新增（AI 占比 高）"),
                   ("s_laglo", "敏感度：甲2 年新增（lag_bldg 低）"), ("s_laghi", "敏感度：甲2 年新增（lag_bldg 高）")):
        CX.cell(r, 1, lab)
        for y in YEARS:
            c = YC[y]
            if k == "adj":
                f = f"=IF(AND({y}>=2027,{y}<=2029),{c}{XT['gw']}-{c}{XT['site']},0)"
            elif k == "site":
                f = "=" + "+".join(f"{c}{XR[p][k]}" for p in CAPP)
            else:   # 甲2 缺值（資本支出最早已知年之前）以 甲 代入
                f = "=" + "+".join(f'IF({c}{XR[p][k]}="",{c}{XR[p]["site"]},{c}{XR[p][k]})' for p in CAPP)
            CX.cell(r, YEARS.index(y) + 4, f).number_format = "0.000"
        XT[k] = r; r += 1
    CX_LAST = r
    CX.column_dimensions["A"].width = 44; CX.column_dimensions["N"].width = 50

    # ================= Deploy：需求情景 × 延遲情景（主路徑＋SiteOnly） =================
    DP["A1"] = ("Deploy — 市場總 GW（年底、年平均、年新增）× 需求情景 A／B／C × 延遲 D0／D1／D2。主路徑：2024–2026＝Pipeline（甲：園區＋實測）；2027–2029 七家改用 甲2"
                "（CapexGW，差額視為年內均勻上線並套用 Pipeline C 節延遲比例）；SiteOnly＝只有 甲。需求情景：2028–29 年新增 × SIF_UnitsCoef_*；2030–32 年新增＝前年年底 × g，"
                "g＝年底 GW 年增率由 2029 值線性收斂到 g_lt（GL01-07）。")
    DP["A2"] = "下游取用 2027 年後數值前看覆蓋率（WIF_Coverage_<yyyy>：有證據支撐的新增 GW ÷ 主路徑新增 GW；< 0.5 時 Checks WARN）"
    hdr(DP, 3, ["項目", "情景", "延遲"] + [str(y) for y in YEARS] + ["說明"])
    r = 4; DR = {}
    def drow(key, lab, s, D, fy, fmt="0.000", fill=None):
        nonlocal r
        DP.cell(r, 1, lab); DP.cell(r, 2, s); DP.cell(r, 3, D)
        for y in YEARS:
            cc = DP.cell(r, YEARS.index(y) + 4, fy(y)); cc.number_format = fmt
            if fill: cc.fill = fill
        DR[key] = r; r += 1
    drow("pipeadd", "Pipeline D0 年新增（淨；甲）", "", "D0", lambda y: f"=Pipeline!{YC[y]}{PA['addD0']}")
    drow("adj", "甲2 調整（CapexGW；2027–2029）", "", "D0", lambda y: f"=CapexGW!{YC[y]}{XT['adj']}")
    drow("stk", "新納入存量", "", "", lambda y: f"=Pipeline!{YC[y]}{AGG['stk']}")
    r += 1
    VAR = (("", "主路徑", True), ("_SO", "SiteOnly（只有 甲）", False))
    for sfx, vlab, use_adj in VAR:
        DP.cell(r, 1, f"■ {vlab}").font = Font(bold=True, size=12); r += 1
        for s in SCEN:
            DP.cell(r, 1, f"需求情景 {s}（{vlab}）").font = BOLD; r += 1
            ra_, rg_, rd_ = r, r + 1, r + 2
            def yeprev(y):
                return f"(Pipeline!{YC[y - 1]}{PA['yeD0']}+SUM($E${rd_}:{YC[y - 1]}{rd_}))"
            def fadd(y, s=s, use_adj=use_adj):
                c = YC[y]
                base_ = f"({c}${DR['pipeadd']}+{c}${DR['adj']})" if use_adj else f"{c}${DR['pipeadd']}"
                if y == 2023: return "=0"
                if y <= 2027: return f"={base_}"
                if y <= 2029: return f"={base_}*{sv('SIF_UnitsCoef_' + s, y)}"
                return f"={yeprev(y)}*{c}{rg_}"
            drow(f"add{s}{sfx}", "年新增（需求情景，D0）", s, "D0", fadd)
            def fg(y):
                if y < 2029: return '=""'
                if y == 2029: return f"=IFERROR({YC[2029]}{ra_}/{yeprev(2029)},g_lt)"
                return f"=${YC[2029]}${rg_}+(g_lt-${YC[2029]}${rg_})*({y}-2029)/3"
            drow(f"g{s}{sfx}", "年底 GW 年增率 g（2029＝情景年新增 ÷ 2028 年底；2030–32 線性收斂到 g_lt）", s, "", fg, "0.0%")
            drow(f"dl{s}{sfx}", "差額（相對 Pipeline D0 年新增；均勻上線）", s, "", lambda y, s=s, sfx=sfx: "=0" if y == 2023 else f"={YC[y]}{DR['add' + s + sfx]}-{YC[y]}${DR['pipeadd']}")
            for D in DLY:
                def fye(z, s=s, D=D, sfx=sfx):
                    return f"=Pipeline!{YC[z]}{PA['ye' + D]}+SUMPRODUCT($E${DR['dl' + s + sfx]}:$M${DR['dl' + s + sfx]},Pipeline!$E${FRAC[D][z]}:$M${FRAC[D][z]})"
                drow(f"ye{s}{D}{sfx}", "年底上線 IT GW", s, D, fye, fill=KEY)
                DP.cell(DR[f"ye{s}{D}{sfx}"], 1).font = BOLD
            for D in DLY:
                drow(f"avg{s}{D}{sfx}", "年平均 IT GW", s, D, lambda y, s=s, D=D, sfx=sfx: '=""' if y == 2023 else f"=({YC[y - 1]}{DR['ye' + s + D + sfx]}+{YC[y]}{DR['ye' + s + D + sfx]})/2")
            for D in DLY:
                drow(f"na{s}{D}{sfx}", "年新增 IT GW（淨，不含新納入存量）", s, D,
                     lambda y, s=s, D=D, sfx=sfx: "=0" if y == 2023 else f"={YC[y]}{DR['ye' + s + D + sfx]}-{YC[y - 1]}{DR['ye' + s + D + sfx]}-{YC[y]}${DR['stk']}")
            for D in ("D1", "D2"):
                drow(f"sno{s}{D}{sfx}", "已出貨未上線 IT GW（相對 D0）", s, D, lambda y, s=s, D=D, sfx=sfx: f"={YC[y]}{DR['ye' + s + 'D0' + sfx]}-{YC[y]}{DR['ye' + s + D + sfx]}")
            r += 1
    DP.cell(r, 1, "覆蓋率（有證據支撐的年新增 ÷ 主路徑年新增；A、D0；MIN 1；主路徑 ≤ 0 時＝1）").font = BOLD
    for y in YEARS:
        c = YC[y]
        DP.cell(r, YEARS.index(y) + 4, f'=IF({y}<=2026,1,IF({c}{DR["naAD0"]}<=0,1,MIN(1,MAX(0,{c}{DR["naAD0_SO"]})/{c}{DR["naAD0"]})))').number_format = "0.00"
    DR["cov"] = r; r += 1
    DP.cell(r, 1, "覆蓋率（含電力證據路徑：SiteOnly 年新增＋T2 納入＝2 的已核准供電合約列；A、D0；MIN 1）").font = BOLD
    for y in YEARS:
        c = YC[y]
        DP.cell(r, YEARS.index(y) + 4, f'=IF({y}<=2026,1,IF({c}{DR["naAD0"]}<=0,1,MIN(1,MAX(0,{c}{DR["naAD0_SO"]}+Pipeline!{c}{AGG["pwr"]})/{c}{DR["naAD0"]})))').number_format = "0.00"
    DR["covp"] = r; r += 2
    DP.cell(r, 1, "對外發布：A-D0、B-D1、C-D2 三條與基準 A-D1（規劃書第 4 節）；Interface 主路徑與 SiteOnly 各 9 組皆輸出").font = BOLD
    DP_LAST = r + 1
    DP.column_dimensions["A"].width = 46
    for col in "BC": DP.column_dimensions[col].width = 6
    DP.column_dimensions["N"].width = 30

    # ================= C0：由下而上 vs 第三方 =================
    C0["A1"] = ("C0 — 市場總 GW：由下而上（Pipeline D0＝Σ 營運者，IT GW）vs 第三方總量（設施 ÷ IF_FacilityGW 換成 IT）。兩讀並列：(a) 由下而上含中國；"
                "(b) 不含中國（假設 Epoch『Other』已含中國買家晶片，C 段 A#116）。第三方只作佐證、不校準；雙方皆 Verified 才判『對得上／偏差』，否則『證據不足（帶內／帶外）』")
    C0["A2"] = ("注意（C 段 6.4）：Epoch 由晶片 × 約 2.5 到設施；本庫晶片 × k_chip2it（基準 1.85）× 1.2＝2.22。倍數未統一前，比值接近 1 不代表由下而上與 Epoch 吻合。"
                "晶片合約（AMD、Broadcom、NVIDIA）不入 C0（Requests R4）")
    hdr(C0, 3, ["項目", "來源鍵", "口徑"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    BUR = {}
    for k, lab, f in (("a", "(a) 由下而上 年底 IT GW（含中國；Pipeline D0）", lambda c: f"=Pipeline!{c}{PA['yeD0']}"),
                      ("alo", "　(a) 低（Players 低列合計）", lambda c: f"=Players!{c}{TOT['lo']}"),
                      ("ahi", "　(a) 高（Players 高列合計）", lambda c: f"=Players!{c}{TOT['hi']}"),
                      ("b", "(b) 由下而上 年底 IT GW（不含中國）", lambda c: f"={c}{BUR['a']}-N(Players!{c}{PR['cn']['tk']})"),
                      ("blo", "　(b) 低", lambda c: f"={c}{BUR['alo']}-N(Players!{c}{PR['cn']['lo']})"),
                      ("bhi", "　(b) 高", lambda c: f"={c}{BUR['ahi']}-N(Players!{c}{PR['cn']['hi']})")):
        C0.cell(r, 1, lab); C0.cell(r, 2, "BU" + k)
        if k in ("a", "b"): C0.cell(r, 1).font = BOLD
        for y in YEARS: C0.cell(r, YEARS.index(y) + 4, f(YC[y])).number_format = "0.000"
        BUR[k] = r; r += 1
    C0.cell(r, 1, "有值玩家數（共 %d 列）" % len(PK))
    for y in YEARS: C0.cell(r, YEARS.index(y) + 4, f"=Players!{YC[y]}{TOT['cnt']}")
    r += 2
    C0R = {}
    for o in seed["thirdparty"]:
        ir = INROW[o["key"]]
        C0.cell(r, 1, o["label"]).font = BOLD; C0.cell(r, 2, o["key"]); C0.cell(r, 3, f"=In!C{ir}")
        r += 1
        C0.cell(r, 1, "　原值（GW）")
        for y in YEARS: C0.cell(r, YEARS.index(y) + 4, f'=IF(In!{IC[y]}{ir}="","",In!{IC[y]}{ir})').font = GREEN
        rv = r; r += 1
        C0.cell(r, 1, "　換算後 IT GW（設施 ÷ IF_FacilityGW）")
        for y in YEARS:
            c = YC[y]
            C0.cell(r, YEARS.index(y) + 4, f'=IF({c}{rv}="","",IF(In!$C${ir}="設施",{c}{rv}/Players!{c}${FACR},{c}{rv}))').number_format = "0.000"
        rc = r; r += 1
        for rd_, lab in (("a", "(a) 含中國"), ("b", "(b) 不含中國")):
            C0.cell(r, 1, f"　比值 {lab}（由下而上 ÷ 第三方；括號＝低–高）")
            for y in YEARS:
                c = YC[y]
                C0.cell(r, YEARS.index(y) + 4, f'=IF(OR({c}{rc}="",N({c}{rc})=0),"",{c}${BUR[rd_]}/{c}{rc})').number_format = "0.00"
            rr_ = r; r += 1
            C0.cell(r, 1, f"　比值區間 {lab}")
            for y in YEARS:
                c = YC[y]
                C0.cell(r, YEARS.index(y) + 4, f'=IF(OR({c}{rc}="",N({c}{rc})=0),"",TEXT({c}${BUR[rd_ + "lo"]}/{c}{rc},"0.00")&"–"&TEXT({c}${BUR[rd_ + "hi"]}/{c}{rc},"0.00"))')
            r += 1
            C0.cell(r, 1, f"　狀態 {lab}")
            for y in YEARS:
                c = YC[y]
                C0.cell(r, YEARS.index(y) + 4,
                        f'=IF({c}{rr_}="","無資料",IF(AND(In!$E${ir}="Verified",FALSE),IF(ABS({c}{rr_}-1)<=tolband,"對得上","偏差"),"證據不足（"&IF(ABS({c}{rr_}-1)<=tolband,"帶內","帶外")&"）"))')
            C0.cell(r, 14, "由下而上為混合標記（Derived），v0.1 一律不判『對得上』")
            C0R[(o["key"], rd_)] = (rr_, r); r += 1
        r += 1
    C0_LAST = r
    C0.column_dimensions["A"].width = 48; C0.column_dimensions["N"].width = 40
    for y in YEARS: C0.column_dimensions[YC[y]].width = 13

    # ================= C1：資本支出對帳 =================
    C1["A1"] = ("C1 — 資本支出對帳：瓦帳推算＝Σ 世代 新增 IT GW × TK IF_CapexIT（每 IT GW；上線年入帳）＋ IF_CapexFacility（提前 lag_bldg 月入帳）；"
                "非 NVIDIA 部分＝持有比 IF_NonNV_Hold_* × VR200 每 GW；＋非 AI 資本支出（In Y3）；曆年按財年結束月換算成財年 vs 個股模型（In Y2）。"
                "第一個已知年的存量不計資本支出；2031–32 廠房缺次年上線量，偏低")
    C1["A2"] = ("MSFT、GOOGL、AMZN、META、ORCL、CRWV、NBIS 2027 起標『同源（不對帳）』：主路徑 2027 起的 GW 由同一份資本支出推得（CapexGW 甲2），再與資本支出比對是循環；"
                "只有 2024–2026（GW 來自園區與實測）做真對帳。本頁比值仍以 甲（SiteOnly）計算，供參考")
    hdr(C1, 3, ["項目", "玩家", "單位"] + [str(y) for y in YEARS] + ["可能原因／說明"])
    r = 4
    C1.cell(r, 1, "廠房提前入帳權重（lag_bldg 月 → l＝月/12；0≤l≤2）").font = BOLD; r += 1
    C1.cell(r, 1, "w0（上線當年）"); C1.cell(r, 4, "=MAX(0,1-lag_bldg/12)").number_format = "0.00"; W0 = f"$D${r}"; r += 1
    C1.cell(r, 1, "w1（前一年）"); C1.cell(r, 4, "=MAX(0,MIN(lag_bldg/12,2-lag_bldg/12))").number_format = "0.00"; W1 = f"$D${r}"; r += 1
    C1.cell(r, 1, "w2（前兩年）"); C1.cell(r, 4, "=MAX(0,MIN(1,lag_bldg/12-1))").number_format = "0.00"; W2 = f"$D${r}"; r += 1
    C1.cell(r, 1, "NVIDIA 每 IT GW IT capex（世代加權，US$B）")
    for y in YEARS:
        c = YC[y]
        C1.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexIT_{g}" for g in GENS)).number_format = "0.00"
    NVIT = r; r += 1
    C1.cell(r, 1, "NVIDIA 每 IT GW 廠房 capex（世代加權，US$B）")
    for y in YEARS:
        c = YC[y]
        C1.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexFacility_{g}" for g in GENS)).number_format = "0.00"
    NVFAC = r; r += 2
    PLAT = {p["key"]: p["nnv_plat"] for p in P}
    REASON = {"msft": "未含融資租賃（10-K 另計）；租用 neocloud 容量（約 1.1 GW）不在資本支出；非 AI 待證據；財年 MW 置於曆年（GL01-04）",
              "googl": "TPU 每 GW 成本以持有比近似；含非 AI（搜尋、YouTube 等）；Epoch 推估 MW",
              "amzn": "資本支出含零售物流（非 AWS）；Trainium 每 GW 成本以持有比近似",
              "orcl": "Stargate 站點多為承租（房東資本支出）；財年 5 月；MW 為模型路徑",
              "crwv": "建物多為房東（Core Scientific、Galaxy）承租；公司資本支出含預付",
              "nbis": "自建＋承租混合；connected 口徑未說明",
              "wyfi": "共識資本支出未取得",
              "meta": "Hyperion、El Paso 為 80/20 合資，資本支出不在 Meta 帳上；含融資租賃本金；租用 CoreWeave、Nebius、Crusoe 不在資本支出",
              "xai": "SpaceX 合併報表未取得；表後發電資本支出另計"}
    C1R = {}
    for p in PK:
        rows = PR[p]; plat = PLAT[p]
        cr = INROW[f"cap_{p}"]; nr = INROW[f"nai_{p}"]
        fye = FYE.get(p, "12")
        C1.cell(r, 1, PN[p]).font = BOLD; C1.cell(r, 2, p); C1.cell(r, 14, REASON.get(p, "個股模型待證據"))
        r += 1
        ra = r
        C1.cell(r, 1, "　IT 資本支出（上線年）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4, f"=Players!{c}{rows['add']}*((1-Players!{c}{rows['nnv']})*{c}${NVIT}+Players!{c}{rows['nnv']}*TK_NonNV_Hold_{plat}*TK_CapexIT_VR200)").number_format = "0.0"
        r += 1; rb = r
        C1.cell(r, 1, "　廠房資本支出（按上線年）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4, f"=Players!{c}{rows['add']}*((1-Players!{c}{rows['nnv']})*{c}${NVFAC}+Players!{c}{rows['nnv']}*TK_NonNV_Hold_{plat}*TK_CapexFacility_VR200)").number_format = "0.0"
        r += 1; rc = r
        C1.cell(r, 1, "　廠房資本支出（入帳曆年，提前 lag_bldg）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            t = [f"{YC[y]}{rb}*{W0}"]
            if y + 1 in YC: t.append(f"{YC[y + 1]}{rb}*{W1}")
            if y + 2 in YC: t.append(f"{YC[y + 2]}{rb}*{W2}")
            C1.cell(r, YEARS.index(y) + 4, "=" + "+".join(t)).number_format = "0.0"
        r += 1; rd = r
        C1.cell(r, 1, "　瓦帳推算 AI 資本支出（曆年）"); C1.cell(r, 3, "US$B")
        for y in YEARS: C1.cell(r, YEARS.index(y) + 4, f"={YC[y]}{ra}+{YC[y]}{rc}").number_format = "0.0"
        r += 1; re_ = r
        C1.cell(r, 1, "　瓦帳推算資本支出（財年；含非 AI）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]
            prev = f"+{YC[y - 1]}{rd}*(12-{fye})/12" if y - 1 in YC else ""
            C1.cell(r, YEARS.index(y) + 4, f"={c}{rd}*{fye}/12{prev}+N(In!{IC[y]}{nr})").number_format = "0.0"
        r += 1; rf = r
        C1.cell(r, 1, "　個股模型資本支出（財年；In Y2）"); C1.cell(r, 3, "US$B")
        for y in YEARS: C1.cell(r, YEARS.index(y) + 4, f'=IF(In!{IC[y]}{cr}="","",In!{IC[y]}{cr})').font = GREEN
        r += 1; rg = r
        C1.cell(r, 1, "　比值（瓦帳 ÷ 個股模型）"); C1.cell(r, 3, "x")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4, f'=IF(OR({c}{rf}="",N({c}{rf})=0),"",{c}{re_}/{c}{rf})').number_format = "0.00"
        r += 1
        C1.cell(r, 1, "　狀態"); C1.cell(r, 3, "—")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4,
                    '="同源（不對帳）"' if (p in ("msft", "googl", "amzn", "meta", "orcl", "crwv", "nbis") and y >= 2027) else
                    f'=IF({c}{rg}="","無個股值",IF({c}{re_}=0,"瓦帳無新增",IF(AND(In!$E${cr}="Verified",FALSE),IF(ABS({c}{rg}-1)<=tolband,"對得上","偏差"),"證據不足（"&IF(ABS({c}{rg}-1)<=tolband,"帶內","帶外")&"）")))')
        C1R[p] = dict(gl=re_, model=rf, ratio=rg, st=r); r += 2
    C1_LAST = r
    C1.column_dimensions["A"].width = 44; C1.column_dimensions["N"].width = 70
    for y in YEARS: C1.column_dimensions[YC[y]].width = 13

    # ================= Compare：瓦帳 vs Siliconomics 成本等值 GW；敏感度 =================
    CP["A1"] = ("Compare — 瓦帳年新增 IT GW（主路徑與 SiteOnly）vs Siliconomics SIF_DeployGW_A/B/C（成本等值 IT GW：AI IT 資本支出 ÷ 每 IT GW IT capex；出貨年）。"
                "兩者口徑不同，只並列、不校準。下半部為敏感度（xAI 宣稱、AI 占比、lag_bldg）")
    hdr(CP, 3, ["項目", "情景", "延遲"] + [str(y) for y in YEARS] + ["說明"])
    r = 4; CPR = {}
    for s in SCEN:
        CP.cell(r, 1, f"Siliconomics 部署 GW（成本等值）"); CP.cell(r, 2, s)
        for y in (2026, 2027, 2028, 2029): CP.cell(r, YEARS.index(y) + 4, f"={sv('SIF_DeployGW_' + s, y)}").number_format = "0.00"
        CP.cell(r, 14, "SIL_In（IT 電力；規劃書 r1）")
        CPR["sil" + s] = r; r += 1
        for sfx, vl in (("", "主路徑"), ("_SO", "SiteOnly")):
            for D in ("D0", "D1"):
                CP.cell(r, 1, f"瓦帳 年新增 IT GW（淨；{vl}）"); CP.cell(r, 2, s); CP.cell(r, 3, D)
                for y in YEARS: CP.cell(r, YEARS.index(y) + 4, f"=Deploy!{YC[y]}{DR['na' + s + D + sfx]}").number_format = "0.00"
                CPR[f"gl{s}{D}{sfx}"] = r; r += 1
            CP.cell(r, 1, f"比值（瓦帳 D0 {vl} ÷ Siliconomics）"); CP.cell(r, 2, s)
            for y in (2026, 2027, 2028, 2029):
                c = YC[y]; CP.cell(r, YEARS.index(y) + 4, f'=IF(N({c}{CPR["sil" + s]})=0,"",{c}{CPR["gl" + s + "D0" + sfx]}/{c}{CPR["sil" + s]})').number_format = "0.00"
            CPR["r" + s + sfx] = r; r += 1
        r += 1
    CP.cell(r, 1, "敏感度（需求 A、D0 年底 IT GW；主路徑）").font = BOLD; r += 1
    CP.cell(r, 1, "主路徑 年底 IT GW"); CP.cell(r, 2, "A"); CP.cell(r, 3, "D0")
    for y in YEARS: CP.cell(r, YEARS.index(y) + 4, f"=Deploy!{YC[y]}{DR['yeAD0']}").number_format = "0.00"
    SB = r; r += 1
    xi = INROW["tgt_xai_27"]
    CP.cell(r, 1, "主路徑 ＋ xAI 宣稱增量（2027 年底 10 GW − 主路徑 xAI；2028 起持平）"); CP.cell(r, 2, "A"); CP.cell(r, 3, "D0")
    for y in YEARS:
        c = YC[y]
        f = f"={c}{SB}" if y < 2027 else f"={c}{SB}+MAX(0,tgt_xai_27-N(Players!{YC[2027]}{PR['xai']['tk']}))"
        CP.cell(r, YEARS.index(y) + 4, f).number_format = "0.00"
    CP.cell(r, 14, f"In 第 {xi} 列（B#46；Interested-party 單一說法，不入主路徑）"); CPR["sxai"] = r; r += 1
    for k, lab in (("s_ailo", "AI 占比 低（ai_sh_lo、ai_amzn_lo）"), ("s_aihi", "AI 占比 高"), ("s_laglo", "lag_bldg 低（12 月）"), ("s_laghi", "lag_bldg 高（24 月）")):
        CP.cell(r, 1, f"主路徑（{lab}）"); CP.cell(r, 2, "A"); CP.cell(r, 3, "D0")
        for y in YEARS:
            c = YC[y]
            if y < 2027:
                f = f"={c}{SB}"
            else:
                zs = [z for z in (2027, 2028, 2029) if z <= y]
                f = f"={c}{SB}+" + "+".join(f"(CapexGW!{YC[z]}{XT[k]}-CapexGW!{YC[z]}{XT['gw']})" for z in zs)
            CP.cell(r, YEARS.index(y) + 4, f).number_format = "0.00"
        CP.cell(r, 14, "2027–2029 甲2 年新增差額累加（2030 起不再傳遞到收斂規則，近似）"); CPR[k] = r; r += 1
    r += 1
    for t_ in ("差額原因（v0.1 第 3 階段）：",
              "1. 出貨年 ≠ 上線年：Siliconomics 是出貨（資本支出）年的成本等值 GW；瓦帳是上線年，建物提前 lag_bldg 入帳。",
              "2. 覆蓋：SiteOnly 只有站點證據（雲端大廠不公布遠期園區；『未找到』不等於『不存在』），2027 起主路徑七家改用 甲2（CapexGW）。其餘玩家（xAI、新雲、主權、其他）仍只有 甲。",
              "3. 成本等值 vs 實體：Siliconomics 以 NVIDIA 每 GW 成本換算 ASIC 支出，ASIC 實體 GW 被低估約 0.3–4.3 GW（2026；規劃書第 6 節第 1 點）。",
              "4. 中國：Siliconomics 顆數含中國出貨（H200 等），瓦帳中國為 Rystad 容量 × AI 占比。"):
        CP.cell(r, 1, t_); r += 1
    CP_LAST = r
    CP.column_dimensions["A"].width = 40; CP.column_dimensions["N"].width = 60

    # ================= FundRaw：F 段融資證據（G 段判定；Funding 頁以 SUMIFS 讀本頁） =================
    fund = json.load(open(FUND, encoding="utf-8"))
    FR["A1"] = ("FundRaw — F_funding.csv 全部列（USD B）＋G 段獨立核對判定與更正（tools/import_funding.py）。採用＝1：G 段判屬實、小誤已更正、§4 重算確認或已改引 6-K；"
                "無法核對與未抽查＝0（只列此頁）。期間桶：2024／2025／2026（最新期）、post（期後）、guide（指引）、other")
    hdr(FR, 3, ["ID", "玩家", "鍵", "期間", "桶", "項目", "值", "單位", "說明", "標記", "來源", "網址", "G 段判定", "G 段更正", "採用"])
    r = 4
    for e in fund["funding"]:
        for j, k in enumerate(("id", "player", "pkey", "period", "bucket", "item", "value", "unit", "detail", "tag", "source", "url", "verdict", "gnote", "use"), 1):
            c = FR.cell(r, j, e[k] if e[k] is not None else "")
            if k == "value" and e["use"]: c.font = BLUE
        if not e["use"]:
            for j in range(1, 16): FR.cell(r, j).font = Font(color="808080")
        r += 1
    FR0, FR1 = 4, r - 1
    for col, w in zip("ABCDEFGHIJKLMNO", (8, 14, 6, 12, 7, 18, 9, 7, 50, 18, 30, 30, 14, 40, 5)):
        FR.column_dimensions[col].width = w
    def fsum(pk, b, items):
        its = items if isinstance(items, (list, tuple)) else [items]
        return "+".join(f'SUMIFS(FundRaw!$G${FR0}:$G${FR1},FundRaw!$C${FR0}:$C${FR1},"{pk}",FundRaw!$E${FR0}:$E${FR1},"{b}",FundRaw!$F${FR0}:$F${FR1},"{it}",FundRaw!$O${FR0}:$O${FR1},1)' for it in its)
    def fcnt(pk, b, items):
        its = items if isinstance(items, (list, tuple)) else [items]
        return "+".join(f'COUNTIFS(FundRaw!$C${FR0}:$C${FR1},"{pk}",FundRaw!$E${FR0}:$E${FR1},"{b}",FundRaw!$F${FR0}:$F${FR1},"{it}",FundRaw!$O${FR0}:$O${FR1},1)' for it in its)

    # ================= Funding：玩家 × 年 資金來源 =================
    FD["A1"] = ("Funding — 玩家 × 期間（2024、2025、2026 最新期；USD B）的資本支出與資金來源。數值以 SUMIFS 由 FundRaw（只取採用＝1）彙總；空白＝未揭露或未採用。"
                "資金來源比例＝支付順序（Decisions GL11-02）：內部現金 → 融資租賃 → 股權 → 債 → 其餘（現金水位／未揭露）；表外與尚未起租另列，不進分母")
    FD["A2"] = "期間標籤見 FundRaw 與 import_funding.py BUCKET（財年：Microsoft 7–6 月、Oracle 6–5 月）；期後事件與指引（post／guide）另列於下方"
    FD["A3"] = ("兩種口徑：【支付順序】假設資本支出先由營業現金流支付、再依序租賃→股權→債，所以營業現金流足以支付時外部資金比例為 0（外部資金被排在營業現金流之後，例如 GOOGL 2025 發債 64.6、"
                "META 2025 發債 30 與 Hyperion 27.3 都不出現）；【總額】＝期間外部資金全額（發債＋定期貸款＋GPU 債＋私募信貸＋新增融資租賃＋表外合資＋股權）÷ 資本支出，與營業現金流 ÷ 資本支出並列，"
                "兩者合計可 > 100%（多出的部分用於回饋、現金水位或期後）。互動頁預設顯示總額口徑")
    FITEMS = [("ocf", "營業現金流", ["ocf"]), ("capex_cash", "現金資本支出", ["capex_cash"]), ("fin_lease", "新增融資租賃", ["fin_lease_new"]),
              ("capex_fl", "含融資租賃資本支出", ["capex_incl_fl"]), ("fcf", "自由現金流", ["fcf"]),
              ("bond", "公司債", ["bond_issue"]), ("term", "定期貸款", ["term_loan"]), ("gpu", "GPU 擔保債", ["gpu_debt"]), ("pc", "私募信貸", ["private_credit"]),
              ("ndc", "淨負債變化", ["net_debt_change"]), ("lnc", "尚未起租租賃承諾（期末存量）", ["lease_not_commenced"]),
              ("jv", "表外合資曝險（當期出資／債）", ["jv_offbalance"]), ("jvx", "合資最大損失曝險（存量）", ["jv_exposure"]),
              ("eq", "股權募資", ["equity_round"]), ("vf", "供應商融資／策略投資", ["vendor_financing"]), ("bb", "買回＋股利", ["buyback_div"]),
              ("talks", "洽談中（不計入）", ["gpu_debt_talks"]), ("refi", "再融資（不重複計入）", ["bridge_refi"])]
    hdr(FD, 4, ["玩家", "鍵", "期間", "期間標籤"] + [x[1] for x in FITEMS] + ["資本支出（分母）", "內部現金", "融資租賃", "股權", "債", "其餘",
                                                     "內部現金 %", "融資租賃 %", "股權 %", "債 %", "其餘 %", "外部募資合計", "債 ÷ 外部", "股 ÷ 外部", "租賃 ÷ 外部",
                                                     "【總額】外部資金", "【總額】外部 ÷ 資本支出", "【總額】營業現金流 ÷ 資本支出", "【總額】債 ÷ 資本支出", "【總額】股 ÷ 資本支出", "【總額】租賃＋表外 ÷ 資本支出"])
    FCOL = {k: get_column_letter(5 + i) for i, (k, _, _) in enumerate(FITEMS)}
    base_ = 5 + len(FITEMS)
    RC = {k: get_column_letter(base_ + i) for i, k in enumerate(("den", "s_int", "s_fl", "s_eq", "s_debt", "s_res", "p_int", "p_fl", "p_eq", "p_debt", "p_res", "ext", "x_debt", "x_eq", "x_fl",
                                                                 "g_ext", "g_extr", "g_ocfr", "g_debtr", "g_eqr", "g_flr"))}
    r = 5; FDR = {}
    FPL = [("msft", "Microsoft"), ("googl", "Alphabet"), ("amzn", "Amazon"), ("meta", "Meta"), ("orcl", "Oracle"), ("crwv", "CoreWeave"), ("nbis", "Nebius"),
           ("xai", "xAI（SpaceX）"), ("oai", "OpenAI"), ("ant", "Anthropic")]
    for pk, pn in FPL:
        for b in ("2024", "2025", "2026"):
            lab = fund["bucket"].get(pn, {}).get(b)
            if not lab: continue
            FD.cell(r, 1, pn); FD.cell(r, 2, pk); FD.cell(r, 3, int(b)); FD.cell(r, 4, lab[0])
            for k, _, its in FITEMS:
                c = FD.cell(r, int(openpyxl.utils.column_index_from_string(FCOL[k])), f'=IF(({fcnt(pk, b, its)})=0,"",{fsum(pk, b, its)})')
                c.number_format = "0.0"
            V = lambda k: f"N({FCOL[k]}{r})"
            den = f'IF({FCOL["capex_fl"]}{r}<>"",{FCOL["capex_fl"]}{r},IF({FCOL["capex_cash"]}{r}<>"",{FCOL["capex_cash"]}{r}+{V("fin_lease")},""))'
            FD[f"{RC['den']}{r}"] = f"={den}"
            D = f"{RC['den']}{r}"
            CC_, CF_ = f"{FCOL['capex_cash']}{r}", f"{FCOL['capex_fl']}{r}"
            FD[f"{RC['s_fl']}{r}"] = f'=IF({D}="","",IF(AND({CF_}<>"",{CC_}<>""),MAX(0,{CF_}-{CC_}),{V("fin_lease")}))'
            FD[f"{RC['s_int']}{r}"] = f'=IF(OR({D}="",{FCOL["ocf"]}{r}=""),"",MAX(0,MIN({V("ocf")},{D}-N({RC["s_fl"]}{r}))))'
            FD[f"{RC['s_eq']}{r}"] = f'=IF({D}="","",MAX(0,MIN({V("eq")},{D}-N({RC["s_int"]}{r})-N({RC["s_fl"]}{r}))))'
            FD[f"{RC['s_debt']}{r}"] = f'=IF({D}="","",MAX(0,MIN({V("bond")}+{V("term")}+{V("gpu")}+{V("pc")},{D}-N({RC["s_int"]}{r})-N({RC["s_fl"]}{r})-N({RC["s_eq"]}{r}))))'
            FD[f"{RC['s_res']}{r}"] = f'=IF({D}="","",{D}-N({RC["s_int"]}{r})-N({RC["s_fl"]}{r})-N({RC["s_eq"]}{r})-N({RC["s_debt"]}{r}))'
            for s_, p_ in (("s_int", "p_int"), ("s_fl", "p_fl"), ("s_eq", "p_eq"), ("s_debt", "p_debt"), ("s_res", "p_res")):
                c = FD.cell(r, openpyxl.utils.column_index_from_string(RC[p_]), f'=IF(OR({D}="",N({D})=0,{RC[s_]}{r}=""),"",{RC[s_]}{r}/{D})'); c.number_format = "0%"; c.fill = KEY
            FD[f"{RC['ext']}{r}"] = f'={V("bond")}+{V("term")}+{V("gpu")}+{V("pc")}+{V("eq")}+{V("fin_lease")}'
            for x_, k_ in (("x_debt", f'{V("bond")}+{V("term")}+{V("gpu")}+{V("pc")}'), ("x_eq", V("eq")), ("x_fl", V("fin_lease"))):
                c = FD.cell(r, openpyxl.utils.column_index_from_string(RC[x_]), f'=IF(N({RC["ext"]}{r})=0,"",({k_})/{RC["ext"]}{r})'); c.number_format = "0%"
            # 總額口徑：期間外部資金全額（不排序）÷ 資本支出；營業現金流 ÷ 資本支出
            FL_ = f'IF(AND({CF_}<>"",{CC_}<>""),MAX(0,{CF_}-{CC_}),{V("fin_lease")})'
            FD[f"{RC['g_ext']}{r}"] = f'={V("bond")}+{V("term")}+{V("gpu")}+{V("pc")}+{FL_}+{V("jv")}+{V("eq")}'
            for k_, num in (("g_extr", f"{RC['g_ext']}{r}"), ("g_ocfr", f'{FCOL["ocf"]}{r}'), ("g_debtr", f'({V("bond")}+{V("term")}+{V("gpu")}+{V("pc")})'),
                            ("g_eqr", V("eq")), ("g_flr", f'({FL_}+{V("jv")})')):
                cond = f'OR({D}="",N({D})=0' + (f',{FCOL["ocf"]}{r}=""' if k_ == "g_ocfr" else "") + ")"
                c = FD.cell(r, openpyxl.utils.column_index_from_string(RC[k_]), f'=IF({cond},"",{num}/{D})'); c.number_format = "0%"
                if k_ in ("g_extr", "g_ocfr"): c.fill = KEY
            FD[f"{RC['g_ext']}{r}"].number_format = "0.0"
            for k in ("den", "s_int", "s_fl", "s_eq", "s_debt", "s_res", "ext"):
                FD[f"{RC[k]}{r}"].number_format = "0.0"
            FDR[(pk, b)] = r; r += 1
    FD_LAST_DATA = r - 1
    r += 1
    FD.cell(r, 1, "期後事件與指引（不進比例；FundRaw 桶 post／guide，採用＝1）").font = BOLD; r += 1
    hdr(FD, r, ["玩家", "期間", "項目", "值", "說明", "ID", "G 段判定"]); r += 1
    for e in fund["funding"]:
        if e["use"] and e["bucket"] in ("post", "guide") and e["pkey"]:
            for j, v in enumerate((e["player"], e["period"], e["item"], e["value"], e["detail"], e["id"], e["verdict"]), 1):
                FD.cell(r, j, v)
            r += 1
    r += 1
    FD.cell(r, 1, "注意：Anthropic 2026 的 SPV 35 已完成（F#190），另 36 為洽談中（F#191，G 段更正）；>71 屬二手轉述、未證實。xAI 2026 過橋 20 為再融資，不與債 25 重複（F#209、F#214）").font = BOLD; r += 2
    FD.cell(r, 1, "循環融資（誰投資誰 → 誰向誰買算力；F.md 第三節 17 條，依 G 段判定標注）").font = BOLD; r += 1
    hdr(FD, r, ["#", "投資方／供應商", "被投資方／客戶", "融資或投資", "回流的算力或採購", "證據（F#）"]); r += 1
    CIRC0 = r
    for i, c_ in enumerate(seed["circular"], 1):
        FD.cell(r, 1, i)
        for j, v in enumerate(c_, 2): FD.cell(r, j, v).alignment = WRAP
        r += 1
    CIRC1 = r - 1
    FD_LAST = r
    FD.column_dimensions["A"].width = 14; FD.column_dimensions["D"].width = 10
    for j in range(5, base_ + 15): FD.column_dimensions[get_column_letter(j)].width = 10

    # ---- C1 廣義口徑對照（寫回 C1 頁；Funding 已建）----
    r = C1_LAST + 1
    C1.cell(r, 1, "C1b. 廣義口徑對照：資本支出（含融資租賃）＋ 尚未起租租賃變化（期末 − 期初）；Microsoft 曆年 2026 起把資料中心租約改列營業租賃，"
                  "報告資本支出會低估實際投入（F#29）。原口徑（上方 C1）保留；本節只在期初、期末尚未起租都有證據時計算").font = BOLD; r += 1
    hdr(C1, r, ["玩家", "鍵", "期間", "期間標籤", "資本支出（含融資租賃）", "尚未起租 期初", "尚未起租 期末", "尚未起租變化", "廣義口徑", "個股模型（財年）", "廣義 ÷ 個股", "說明"]); r += 1
    C1B0 = r
    for pk in ("msft", "googl", "amzn", "meta", "orcl", "crwv", "nbis"):
        for b in ("2024", "2025", "2026"):
            fr = FDR.get((pk, b))
            if not fr: continue
            prv = FDR.get((pk, str(int(b) - 1)))
            C1.cell(r, 1, PN[pk]); C1.cell(r, 2, pk); C1.cell(r, 3, int(b)); C1.cell(r, 4, f"=Funding!D{fr}").font = GREEN
            C1.cell(r, 5, f"=Funding!{RC['den']}{fr}").font = GREEN
            C1.cell(r, 6, f"=Funding!{FCOL['lnc']}{prv}" if prv else '=""').font = GREEN
            C1.cell(r, 7, f"=Funding!{FCOL['lnc']}{fr}").font = GREEN
            C1.cell(r, 8, f'=IF(OR(F{r}="",G{r}=""),"",G{r}-F{r})').number_format = "0.0"
            C1.cell(r, 9, f'=IF(OR(E{r}="",H{r}=""),"",E{r}+H{r})').number_format = "0.0"
            cr = INROW[f"cap_{pk}"]
            C1.cell(r, 10, f'=IF(In!{IC[int(b)]}{cr}="","",In!{IC[int(b)]}{cr})').font = GREEN
            C1.cell(r, 11, f'=IF(OR(I{r}="",J{r}="",LEFT(D{r},2)="20"),"",I{r}/J{r})').number_format = "0.00"
            C1.cell(r, 12, "期間非全年（H1／9M）時不算比值；尚未起租為承諾（未來 1–20 年起租），廣義口徑是上限讀法")
            r += 1
    C1_LAST = r
    C1.column_dimensions["L"].width = 50

    # ================= Facility：每 IT GW 廠房資本支出拆分與市場層 =================
    FC["A1"] = ("Facility — 每 IT GW 廠房資本支出（TK IF_CapexFacility，世代加權；VR200 12.67 $B）拆分（P_facility_power.md 建議表，G 段核對）；"
                "備援發電機改為倍數 × 單價（基準 1.6 × 1.0，Decisions GL11-04），差額從『其他專業服務』扣；殘差列『其他／未拆』使合計＝TK。"
                "自建發電另列（不在 12.67 內）：每 IT GW 1.44 GW × 單價，乘自建比例。市場金額＝主路徑（A、D0）年新增 IT GW × 每 GW 各項")
    hdr(FC, 3, ["項目", "鍵", "單位"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    FC.cell(r, 1, "每 IT GW 廠房資本支出（TK，世代加權）").font = BOLD
    for y in YEARS:
        c = YC[y]
        FC.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexFacility_{g}" for g in GENS)).number_format = "0.00"
    FTK = r; r += 1
    FACI = [("dist", "配電", "=fac_dist"), ("gen", "備援發電機（倍數 × 單價）", "=fac_gen_mult*fac_gen_px"), ("ups", "UPS 與電池", "=fac_ups"),
            ("cool", "冷卻", "=fac_cool"), ("bldg", "建築（殼體與結構）", "=fac_bldg"), ("labor", "施工人工與工程設計", "=fac_labor"),
            ("prof", "其他專業服務（扣 genset 調整差額）", "=fac_prof0-(fac_gen_mult*fac_gen_px-fac_gen_px0)"), ("net", "網路、消防、安防、BMS", "=fac_net"),
            ("land", "土地", "=fac_land")]
    FCR = {}
    for k, lab, f in FACI:
        FC.cell(r, 1, "　" + lab); FC.cell(r, 2, f"fac_{k}"); FC.cell(r, 3, "$B/IT GW")
        for y in YEARS:
            FC.cell(r, YEARS.index(y) + 4, f).number_format = "0.00"
        FCR[k] = r; r += 1
    FC.cell(r, 1, "　其他／未拆（殘差＝TK − Σ 上列）"); FC.cell(r, 2, "fac_resid"); FC.cell(r, 3, "$B/IT GW")
    for y in YEARS:
        c = YC[y]
        FC.cell(r, YEARS.index(y) + 4, f"={c}{FTK}-SUM({c}{FCR['dist']}:{c}{FCR['land']})").number_format = "0.00"
    FCR["resid"] = r; r += 1
    FC.cell(r, 1, "　拆分合計（須＝TK）"); FC.cell(r, 3, "$B/IT GW")
    for y in YEARS:
        c = YC[y]; FC.cell(r, YEARS.index(y) + 4, f"=SUM({c}{FCR['dist']}:{c}{FCR['resid']})").number_format = "0.00"
    FCR["sum"] = r; r += 1
    FC.cell(r, 1, "　自建發電（另計，不在 TK 廠房內；僅有自建的園區）"); FC.cell(r, 2, "fac_self"); FC.cell(r, 3, "$B/IT GW")
    for y in YEARS: FC.cell(r, YEARS.index(y) + 4, "=self_gw*self_px").number_format = "0.00"
    FCR["self_unit"] = r; r += 2
    FC.cell(r, 1, "實物量（每 IT GW）").font = BOLD; r += 1
    for k, lab, f, unit in (("q_fac", "設施電力", None, "GW"), ("q_gen", "備援發電機", "=fac_gen_mult", "GW"), ("q_self", "自建發電（有自建者）", "=self_gw", "GW"),
                            ("q_xfmr", "主變壓器（含一台退出備援；P#50 ModulEdge、ComEd 例）", "=xfmr_gva", "GVA")):
        FC.cell(r, 1, "　" + lab); FC.cell(r, 3, unit)
        for y in YEARS:
            FC.cell(r, YEARS.index(y) + 4, f if f else f"=Players!{YC[y]}{FACR}").number_format = "0.00"
        FCR[k] = r; r += 1
    r += 1
    FC.cell(r, 1, "市場層（主路徑 A、D0 年新增 IT GW × 每 GW 各項；US$B）").font = BOLD; r += 1
    FC.cell(r, 1, "　年新增 IT GW（主路徑 A、D0）"); FC.cell(r, 3, "GW")
    for y in YEARS: FC.cell(r, YEARS.index(y) + 4, f"=MAX(0,Deploy!{YC[y]}{DR['naAD0']})").number_format = "0.00"
    FADD = r; r += 1
    FMK = {}
    for k, lab, _ in FACI + [("resid", "其他／未拆", None)]:
        FC.cell(r, 1, f"　市場：{lab}"); FC.cell(r, 2, f"mkt_{k}"); FC.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]; FC.cell(r, YEARS.index(y) + 4, f"={c}${FADD}*{c}{FCR[k]}").number_format = "0.0"
        FMK[k] = r; r += 1
    FC.cell(r, 1, "　市場：廠房合計（TK）"); FC.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; FC.cell(r, YEARS.index(y) + 4, f"=SUM({c}{FMK['dist']}:{c}{FMK['resid']})").number_format = "0.0"
    FMK["sum"] = r; r += 1
    FC.cell(r, 1, "　市場：自建發電（× 自建比例 self_sh）"); FC.cell(r, 2, "mkt_self"); FC.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; FC.cell(r, YEARS.index(y) + 4, f"={c}${FADD}*self_sh*{c}{FCR['self_unit']}").number_format = "0.0"
    FMK["self"] = r; r += 1
    FC.cell(r, 1, "　市場：備援發電機裝置量（GW）"); FC.cell(r, 3, "GW")
    for y in YEARS:
        c = YC[y]; FC.cell(r, YEARS.index(y) + 4, f"={c}${FADD}*fac_gen_mult").number_format = "0.0"
    FMK["q_gen"] = r; r += 1
    FC.cell(r, 1, "　市場：自建發電容量（GW）"); FC.cell(r, 3, "GW")
    for y in YEARS:
        c = YC[y]; FC.cell(r, YEARS.index(y) + 4, f"={c}${FADD}*self_sh*self_gw").number_format = "0.0"
    FMK["q_self"] = r; r += 2
    FC.cell(r, 1, "延遲佐證：美國表後自建發電 規劃 btm_plan 90 GW／運轉 btm_run 2 GW（P#39，Cleanview；G 段屬實）→ 運轉 ÷ 規劃").font = BOLD
    FC.cell(r, 4, "=btm_run/btm_plan").number_format = "0.0%"; r += 1
    FC_LAST = r
    FC.column_dimensions["A"].width = 46; FC.column_dimensions["N"].width = 40

    # ================= SegIn：AI半導體 次產業比例鏡像 =================
    import csv as _csv
    SG["A1"] = (f"SegIn — AI半導體 次產業『GB300 對照 ÷ IT』比例（只讀鏡像；權威在 AI半導體）：markets/invest_map/data/segments_ai_semis_v0.2.3.csv（commit {SEG_COMMIT}）。"
                "比例為 NVIDIA 平台橋接表（合計 100%），Spend 頁套到全部 IT capex（ASIC 為近似）")
    hdr(SG, 3, ["次產業", "鍵", "GB300 對照 ÷ IT", "2027 金額 ÷ AI IT capex（參考）", "金額標記", "來源欄"])
    SEGK = {"運算": "compute", "記憶體": "memory", "網通": "network", "代工與封測": "foundry", "載板與 PCB": "substrate", "電源與散熱": "power_thermal",
            "伺服器組裝": "server", "儲存與管理": "storage"}
    r = 4; SGR = {}
    for row_ in _csv.DictReader(open(SEG, encoding="utf-8-sig")):
        nm_ = row_["次產業"]
        if nm_ not in SEGK: continue
        v = row_["GB300 對照 ÷ IT"].strip().rstrip("%")
        SG.cell(r, 1, nm_); SG.cell(r, 2, SEGK[nm_])
        c = SG.cell(r, 3, float(v) / 100); c.font = GREEN; c.number_format = "0.0%"
        SG.cell(r, 4, row_["2027 金額 ÷ AI IT capex"]); SG.cell(r, 5, row_["金額標記"]); SG.cell(r, 6, "GB300 對照 ÷ IT")
        name(wb, f"SEG_{SEGK[nm_]}", f"SegIn!$C${r}"); SGR[SEGK[nm_]] = r; r += 1
    SG.cell(r, 1, "合計（須＝100%）").font = BOLD
    SG.cell(r, 3, f"=SUM(C4:C{r - 1})").number_format = "0.0%"; SGSUM = r
    SG.column_dimensions["A"].width = 16; SG.column_dimensions["D"].width = 28

    # ================= Spend：年 × 項目 市場金額 =================
    SP["A1"] = ("Spend — 年 × 項目的 AI 資料中心市場金額（US$B）＝主路徑（A、D0）年新增 IT GW ×（每 IT GW IT capex × SegIn 次產業比例）＋ Facility 各項 ＋ 自建發電。"
                "每 IT GW IT capex＝(1 − 市場 NonNV 比) × NVIDIA 世代加權 IF_CapexIT ＋ NonNV 比 × spend_nnv_hold × VR200 IF_CapexIT（Decisions GL11-05）。上線年口徑（非出貨年）")
    hdr(SP, 3, ["項目", "鍵", "單位"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    SP.cell(r, 1, "年新增 IT GW（主路徑 A、D0）"); SP.cell(r, 3, "GW")
    for y in YEARS: SP.cell(r, YEARS.index(y) + 4, f"=Facility!{YC[y]}{FADD}").number_format = "0.00"
    SADD = r; r += 1
    SP.cell(r, 1, "每 IT GW IT capex（US$B）"); SP.cell(r, 3, "$B/IT GW")
    for y in YEARS:
        c = YC[y]
        nv = "+".join(f"Players!{c}{MIXR[g]}*TK_CapexIT_{g}" for g in GENS)
        SP.cell(r, YEARS.index(y) + 4, f"=(1-Players!{c}{NNVR})*({nv})+Players!{c}{NNVR}*spend_nnv_hold*TK_CapexIT_VR200").number_format = "0.00"
    SIT = r; r += 2
    SPR = {}
    SP.cell(r, 1, "A. IT 設備（依 AI半導體 次產業比例）").font = BOLD; r += 1
    SEGL = {"compute": "運算", "memory": "記憶體", "network": "網通", "foundry": "代工封測", "substrate": "載板與 PCB", "power_thermal": "電源散熱（機架內）",
            "server": "伺服器組裝", "storage": "儲存與管理"}
    for k in ("compute", "memory", "network", "foundry", "substrate", "power_thermal", "server", "storage"):
        SP.cell(r, 1, "　" + SEGL[k]); SP.cell(r, 2, f"it_{k}"); SP.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]; SP.cell(r, YEARS.index(y) + 4, f"={c}${SADD}*{c}${SIT}*SEG_{k}").number_format = "0.0"
        SPR["it_" + k] = r; r += 1
    SP.cell(r, 1, "　IT 設備小計"); SP.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; SP.cell(r, YEARS.index(y) + 4, f"=SUM({c}{SPR['it_compute']}:{c}{SPR['it_storage']})").number_format = "0.0"
    SPR["it_sum"] = r; r += 1
    SP.cell(r, 1, "B. 廠房（Facility 頁）").font = BOLD; r += 1
    for k, lab, _ in FACI + [("resid", "其他／未拆", None)]:
        SP.cell(r, 1, "　" + lab); SP.cell(r, 2, f"fac_{k}"); SP.cell(r, 3, "US$B")
        for y in YEARS: SP.cell(r, YEARS.index(y) + 4, f"=Facility!{YC[y]}{FMK[k]}").font = GREEN
        SPR["fac_" + k] = r; r += 1
    SP.cell(r, 1, "　廠房小計"); SP.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; SP.cell(r, YEARS.index(y) + 4, f"=SUM({c}{SPR['fac_dist']}:{c}{SPR['fac_resid']})").number_format = "0.0"
    SPR["fac_sum"] = r; r += 1
    SP.cell(r, 1, "C. 自建發電（不在 TK 內；× 自建比例）"); SP.cell(r, 2, "self_gen"); SP.cell(r, 3, "US$B")
    for y in YEARS: SP.cell(r, YEARS.index(y) + 4, f"=Facility!{YC[y]}{FMK['self']}").font = GREEN
    SPR["self_gen"] = r; r += 1
    SP.cell(r, 1, "合計（IT＋廠房＋自建發電）").font = BOLD; SP.cell(r, 3, "US$B")
    for y in YEARS:
        c = YC[y]; cc = SP.cell(r, YEARS.index(y) + 4, f"={c}{SPR['it_sum']}+{c}{SPR['fac_sum']}+{c}{SPR['self_gen']}"); cc.number_format = "0.0"; cc.fill = KEY
    SPR["total"] = r; r += 1
    SP.cell(r, 1, "對照：Siliconomics AI IT 資本支出（情境 A；出貨年）"); SP.cell(r, 3, "US$B")
    for y in (2026, 2027, 2028, 2029): SP.cell(r, YEARS.index(y) + 4, f"={sv('SIF_AIITCapex_A', y)}").number_format = "0.0"
    SPR["sil"] = r; r += 1
    SP_LAST = r
    SP.column_dimensions["A"].width = 44

    # ================= Gap：資金缺口（v0.1.7；規劃 docs/plan/20261031_主要玩家資金缺口_規劃.md、工作單 20261031_Wattonomics_v0.1.7_資金缺口.md） =================
    GA["A1"] = ("Gap — 資金缺口：需要的錢＝新增 IT GW × 每 IT GW 建置成本（TK IF_CapexTotal＝IT＋廠房，與 CapexGW／C1／Spend 同一組儲存格；廠房依 lag_bldg 提前入帳），"
                "每 GW 只在擁有者計一次；錢從哪來＝主要玩家全公司營業現金流（本業＋AI；不扣買回與股利，Andy 2026-10-31 04:20）＋已觀測外部融資（2024–2025，Funding 總額口徑）。"
                "2026 起已到位資金只算 2025 年底現金與尚未起租租賃、表外合資存量，只進累計（判斷題 G4）。US$B；曆年")
    GA["A2"] = ("主要玩家＝11 家合計（Microsoft、Alphabet、Amazon、Meta、Oracle、CoreWeave、Nebius、WhiteFiber、SpaceX／xAI、OpenAI、Anthropic）＋智譜只列。"
                "營業現金流固定為個股模型基準情境，不隨需求 A／B／C 變（G5）；2028 起覆蓋率 < 0.35 屬情景（G6）；個股模型沒有的年份標『無模型值』、不外推。"
                "OpenAI／Stargate 自建 GW 由 SoftBank、SB Energy 等 SPV 擁有，依規劃書 §3.3 歸『市場其餘』，不計 OpenAI 資本支出（Decisions GL01-31）")
    hdr(GA, 3, ["項目", "鍵", "單位"] + [str(y) for y in YEARS] + ["說明"])
    OUTG = list(range(2024, 2031))
    GROW = {}
    r = 4

    def gline(key, lab, unit, fy, years=YEARS, fmt="0.0", note=None, bold=False, fill=None):
        nonlocal r
        GA.cell(r, 1, lab); GA.cell(r, 2, key); GA.cell(r, 3, unit)
        if bold: GA.cell(r, 1).font = BOLD
        for y in years:
            f = fy(y)
            if f is None: continue
            cc = GA.cell(r, YEARS.index(y) + 4, f); cc.number_format = fmt
            if fill: cc.fill = fill
        if note: GA.cell(r, 14, note)
        GROW[key] = r; r += 1
        return r - 1

    def gsec(t):
        nonlocal r
        r += 1; GA.cell(r, 1, t).font = Font(bold=True, size=12); r += 1

    G = lambda key, y: f"{YC[y]}{GROW[key]}"
    # ---- A. 每 IT GW 建置成本（市場） ----
    gsec("A. 每 IT GW 建置成本（市場；US$B／IT GW；IT＝Spend 頁同一公式，廠房＝CapexGW 同一公式套市場 NonNV 比）")
    GA.cell(r, 1, "廠房提前入帳權重 w0／w1／w2（引用 CapexGW；lag_bldg）")
    for j, w in enumerate(XW): GA.cell(r, 4 + j, f"={w}").number_format = "0.00"
    GROW["w"] = r; r += 1
    GW0, GW1, GW2 = (f"Gap!$D${GROW['w']}", f"Gap!$E${GROW['w']}", f"Gap!$F${GROW['w']}")
    gline("nnv", "市場新增 GW 非 NVIDIA 占比（Players）", "%", lambda y: f"=Players!{YC[y]}{NNVR}", fmt="0.0%")
    gline("pit", "每 IT GW IT capex（＝Spend 頁）", "$B/IT GW", lambda y: f"=Spend!{YC[y]}{SIT}", fmt="0.00",
          note="(1 − NonNV) × NVIDIA 世代加權 IF_CapexIT ＋ NonNV × spend_nnv_hold × VR200 IF_CapexIT（GL11-05）")
    gline("pfac", "每 IT GW 廠房 capex", "$B/IT GW",
          lambda y: f"=(1-{G('nnv', y)})*CapexGW!{YC[y]}{XFAC}+{G('nnv', y)}*spend_nnv_hold*TK_CapexFacility_VR200", fmt="0.00",
          note="(1 − NonNV) × NVIDIA 世代加權 IF_CapexFacility ＋ NonNV × spend_nnv_hold × VR200 IF_CapexFacility（CapexGW 同式，市場 NonNV）")
    gline("pgw", "每 IT GW 建置成本合計（IT＋廠房）", "$B/IT GW", lambda y: f"={G('pit', y)}+{G('pfac', y)}", fmt="0.00", bold=True, fill=KEY)
    GA.cell(r, 1, "　項目展開（IT 依 SegIn 次產業比例；廠房依 Facility 頁拆分比例 × 本頁廠房每 GW）").font = BOLD; r += 1
    GITEMS = [("it_" + k, "IT：" + v, f"*SEG_{k}") for k, v in (("compute", "運算晶片"), ("memory", "記憶體"), ("network", "網通與光模組"), ("foundry", "代工封測"),
                                                                ("substrate", "載板與 PCB"), ("power_thermal", "機架內電源散熱"), ("server", "伺服器組裝"), ("storage", "儲存與管理"))]
    GFAC = [("fac_" + k, "廠房：" + lab, k) for k, lab in (("dist", "配電"), ("gen", "備援發電機"), ("ups", "UPS 與電池"), ("cool", "冷卻"), ("bldg", "建築"),
                                                         ("labor", "施工人工與設計"), ("prof", "其他專業服務"), ("net", "網路、消防、安防"), ("land", "土地"), ("resid", "其他／未拆"))]
    for k, lab, m in GITEMS:
        gline("i_" + k, "　" + lab, "$B/IT GW", lambda y, m=m: f"={G('pit', y)}{m}", fmt="0.00")
    for k, lab, fk in GFAC:
        gline("i_" + k, "　" + lab, "$B/IT GW", lambda y, fk=fk: f"={G('pfac', y)}*Facility!{YC[y]}{FCR[fk]}/Facility!{YC[y]}{FTK}", fmt="0.00")
    gline("isum", "　項目合計（須＝每 GW 建置成本）", "$B/IT GW", lambda y: f"=SUM({YC[y]}{GROW['i_' + GITEMS[0][0]]}:{YC[y]}{GROW['i_' + GFAC[-1][0]]})", fmt="0.00")
    gline("iself", "　備忘：自建發電（Spend 另計，不在 IF_CapexTotal，不進需要）", "$B/IT GW", lambda y: f"=Facility!{YC[y]}{FCR['self_unit']}*self_sh", fmt="0.00")
    gline("cov", "覆蓋率（WIF_Coverage；A、D0）", "比", lambda y: f"=Deploy!{YC[y]}{DR['cov']}", fmt="0.00")
    gline("covflag", "年份性質", "—", lambda y: f'=IF({G("cov", y)}<0.35,"情景（覆蓋率 < 0.35）","")', years=OUTG, fmt="General")
    GA.cell(r, 1, "對照：年化持有成本（TK AL_HoldGen／IF_HoldGen 世代加權 × 年平均 GW；G1 對照列）")
    GA.cell(r, 4, "待 TK：Tokenomics 快照 v21.7 未含年化持有成本名稱（Requests R6）；本版不列數值，不自行估算"); GROW["hold"] = r; r += 1

    # ---- B. 玩家層 ----
    gsec("B. 玩家層（11 家＋智譜只列；曆年；US$B）。AI 資本支出＝玩家年新增 IT GW（主路徑：七家 2027 起甲2，其餘甲）× 玩家每 GW（CapexGW／C1 同式），廠房提前 lag_bldg")
    CAPP7 = set(XR)
    FULL = {}
    for (pk, b), fr in FDR.items():
        pn = dict(FPL)[pk]
        lab_ = fund["bucket"].get(pn, {}).get(b, [""])[0]
        FULL[(pk, int(b))] = (fr, str(lab_).startswith(("CY", "FY")))
    GPR = {}
    GAPREF = {}   # 缺口列引用（G-K5 公式文字檢查用）

    def fyexpr(p, inkey, y, fund_col=None, scale=""):
        """財年值：In Y9（模型）優先；否則 Funding 全年觀測（只用期間標籤為 CY／FY 的全年列）"""
        ir = INROW[inkey]; inref = f"In!{IC[y]}{ir}"
        fb = FULL.get((p, y)) if fund_col else None
        if fb and fb[1]:
            fc = f"Funding!{fund_col}{fb[0]}"
            return f'IF({inref}<>"",{inref}{scale},IF({fc}<>"",{fc},""))'
        return f'IF({inref}<>"",{inref}{scale},"")'

    def cyexpr(rowkey, y, fye):
        """曆年＝財年 × fye/12 ＋ 次財年 ×（12 − fye）/12（CapexGW、C1 同法）；缺任一 → 空白"""
        a_ = f"{YC[y]}{GPR[p_cur][rowkey]}"
        if fye == 12: return f"={a_}"
        b_ = f"{YC[y + 1]}{GPR[p_cur][rowkey]}"
        return f'=IF(OR({a_}="",{b_}=""),"",{a_}*{fye}/12+{b_}*{12 - fye}/12)'
    for p in GP:
        p_cur = p
        fye = int(GSEED["players"].get(p, {}).get("fye", 12))
        zs = "/(10*gfx_cny)" if p == "zhipu" else ""
        GA.cell(r, 1, f"{GPN[p]}" + ("（只列，不進合計；RMB 億 ÷ 10 ÷ gfx_cny）" if p == "zhipu" else "")).font = BOLD; GA.cell(r, 2, p)
        GA.cell(r, 14, f"財年結束月 {fye}；來源 {GSRC[p]}（In Y9）"); r += 1
        rows = {}; GPR[p] = rows

        def pl(k, lab, unit, fy, years=YEARS, fmt="0.0", note=None, fill=None):
            nonlocal r
            GA.cell(r, 1, "　" + lab); GA.cell(r, 2, f"{p}_{k}"); GA.cell(r, 3, unit)
            for y in years:
                f = fy(y)
                if f is None: continue
                cc = GA.cell(r, YEARS.index(y) + 4, f); cc.number_format = fmt
                if fill: cc.fill = fill
            if note: GA.cell(r, 14, note)
            rows[k] = r; r += 1
        if p != "zhipu":
            if p in CAPP7:
                pl("add", "年新增 IT GW（主路徑；擁有者）", "GW", lambda y: f"=CapexGW!{YC[y]}{XR[p]['main']}", fmt="0.000", note="CapexGW 主路徑（2027 起＝甲2）")
                pl("pit", "每 IT GW IT capex（CapexGW）", "$B/IT GW", lambda y: f"=CapexGW!{YC[y]}{XR[p]['pit']}", fmt="0.00")
                pl("pfac", "每 IT GW 廠房 capex（CapexGW）", "$B/IT GW", lambda y: f"=CapexGW!{YC[y]}{XR[p]['pfac']}", fmt="0.00")
            else:
                if p == "oai":
                    pl("spv", "備忘：OpenAI／Stargate 自建 年新增 IT GW（SPV 擁有，歸市場其餘）", "GW", lambda y: f"=Players!{YC[y]}{PR[p]['add']}", fmt="0.000",
                       note="規劃書 §3.3：Stargate SoftBank、SB Energy 等 SPV 視同市場其餘的擁有者；OpenAI 只承擔租金（已在營業現金流）")
                    pl("add", "年新增 IT GW（擁有者口徑）", "GW", lambda y: "=0", fmt="0.000", note="SPV 擁有 → 0（GL01-31）")
                else:
                    pl("add", "年新增 IT GW（Players 甲；擁有者）", "GW", lambda y: f"=Players!{YC[y]}{PR[p]['add']}", fmt="0.000")
                nv = lambda y: f"Players!{YC[y]}{PR[p]['nnv']}"
                pl("pit", "每 IT GW IT capex（C1 同式）", "$B/IT GW", lambda y: f"=(1-{nv(y)})*CapexGW!{YC[y]}{XIT}+{nv(y)}*TK_NonNV_Hold_{PLAT[p]}*TK_CapexIT_VR200", fmt="0.00")
                pl("pfac", "每 IT GW 廠房 capex（C1 同式）", "$B/IT GW", lambda y: f"=(1-{nv(y)})*CapexGW!{YC[y]}{XFAC}+{nv(y)}*TK_NonNV_Hold_{PLAT[p]}*TK_CapexFacility_VR200", fmt="0.00")

            def aif(y):
                c, t = YC[y], []
                t.append(f"{c}{rows['add']}*({c}{rows['pit']}+{GW0}*{c}{rows['pfac']})")
                for k_, w_ in ((1, GW1), (2, GW2)):
                    if y + k_ in YC: t.append(f"{w_}*{YC[y + k_]}{rows['add']}*{YC[y + k_]}{rows['pfac']}")
                return "=" + "+".join(t)
            pl("ai", "AI 資料中心資本支出（新增 GW × 每 GW；廠房提前）", "US$B", aif, years=OUTG, fill=KEY)
        else:
            pl("ai", "AI 資料中心資本支出", "US$B", lambda y: "=0", years=OUTG, note="智譜不是 Wattonomics 營運者（容量在『中國合計』）；自有算力資本支出基準 0（模型 D11r）")
        # 非 AI、模型資本支出、營業現金流、買回股利（財年 → 曆年）
        pl("nai_fy", "非 AI 資本支出（財年；模型）", "US$B", lambda y: "=" + fyexpr(p, f"gnai_{p}", y, None, zs))
        pl("nai_cy", "非 AI 資本支出（曆年；空白＝無模型值）", "US$B", lambda y: cyexpr("nai_fy", y, fye), years=OUTG)
        pl("nai", "非 AI 資本支出（曆年；無模型值以 0 計，G2）", "US$B", lambda y: f"=N({YC[y]}{rows['nai_cy']})", years=OUTG)
        pl("cap_fy", "資本支出合計（財年；模型；對照）", "US$B", lambda y: "=" + fyexpr(p, f"gcap_{p}", y, None, zs))
        pl("cap", "資本支出合計（曆年；模型；對照）", "US$B", lambda y: cyexpr("cap_fy", y, fye), years=OUTG)
        pl("cai_fy", "AI 資本支出（財年；模型；對照）", "US$B", lambda y: "=" + fyexpr(p, f"gcapai_{p}", y, None, zs))
        pl("cai", "對照：個股模型 AI 資本支出（曆年；空白＝無模型值）", "US$B", lambda y: cyexpr("cai_fy", y, fye), years=OUTG)
        if p in CAPP7:
            pl("capx", "對照：CapexGW 曆年 AI 資本支出（In Y2 × AI 占比；C1 對照）", "US$B", lambda y: f"=CapexGW!{YC[y]}{XR[p]['ai']}", years=OUTG)
        ocol = FCOL["ocf"] if p in dict(FPL) else None
        pl("ocf_fy", "營業現金流（財年；模型，缺則 Funding 全年觀測）", "US$B", lambda y: "=" + fyexpr(p, f"gocf_{p}", y, ocol, zs))
        pl("ocf_src", "　營業現金流來源（財年）", "—", lambda y: (f'=IF(In!{IC[y]}{INROW["gocf_" + p]}<>"","模型",IF({YC[y]}{rows["ocf_fy"]}<>"","Funding 觀測","無模型值"))'), fmt="General")
        pl("ocf", "營業現金流（曆年；全公司；空白＝無模型值）", "US$B", lambda y: cyexpr("ocf_fy", y, fye), years=OUTG, fill=KEY)
        bcol = FCOL["bb"] if p in dict(FPL) else None
        pl("bb_fy", "買回＋股利（財年；備忘）", "US$B", lambda y: "=" + fyexpr(p, f"gbb_{p}", y, bcol, zs))
        pl("bb", "備忘：可削減緩衝＝買回＋股利（曆年；不進缺口）", "US$B", lambda y: cyexpr("bb_fy", y, fye), years=OUTG)
        pl("need", "需要合計（AI＋非 AI）", "US$B", lambda y: f"={YC[y]}{rows['ai']}+{YC[y]}{rows['nai']}", years=OUTG)
        pl("self", "自籌比＝營業現金流 ÷（AI＋非 AI）", "x", lambda y: f'=IF(OR({YC[y]}{rows["ocf"]}="",{YC[y]}{rows["need"]}<=0),"",{YC[y]}{rows["ocf"]}/{YC[y]}{rows["need"]})', years=OUTG, fmt="0.00")
        pl("gap", "缺口＝（AI＋非 AI）− 營業現金流（正＝需外部資金）", "US$B", lambda y: f'=IF({YC[y]}{rows["ocf"]}="","",{YC[y]}{rows["need"]}-{YC[y]}{rows["ocf"]})', years=OUTG, fill=KEY)
        GAPREF[p] = (rows["gap"], rows["self"], rows["bb"], rows["bb_fy"])
        # 已觀測融資（2024–2025；Funding 總額口徑；財年公司換曆年）
        if p in dict(FPL):
            def fext(y, piece):
                def one(yy):
                    fb = FULL.get((p, yy))
                    if not fb or not fb[1]: return None
                    fr = fb[0]
                    debt = f"N(Funding!{FCOL['bond']}{fr})+N(Funding!{FCOL['term']}{fr})+N(Funding!{FCOL['gpu']}{fr})+N(Funding!{FCOL['pc']}{fr})"
                    return {"all": f"N(Funding!{RC['g_ext']}{fr})", "debt": f"({debt})", "eq": f"N(Funding!{FCOL['eq']}{fr})",
                            "lj": f"(N(Funding!{RC['g_ext']}{fr})-({debt})-N(Funding!{FCOL['eq']}{fr}))"}[piece]
                a_ = one(y)
                if fye == 12: return f"={a_ or 0}"
                b_ = one(y + 1)
                return f"={a_ or 0}*{fye}/12+{b_ or 0}*{12 - fye}/12"
            for k_, lab_ in (("ext", "已觀測外部融資（Funding 總額口徑）"), ("ext_debt", "　其中：債（公司債、定期貸款、GPU 擔保、私募信貸）"),
                             ("ext_eq", "　其中：股權"), ("ext_lj", "　其中：新增融資租賃＋表外合資")):
                pl(k_, lab_ + ("；財年按月數換曆年" if fye != 12 else ""), "US$B", lambda y, k_=k_: fext(y, {"ext": "all", "ext_debt": "debt", "ext_eq": "eq", "ext_lj": "lj"}[k_]),
                   years=[2024, 2025], fill=KEY if k_ == "ext" else None)
            pl("imp", "推算尚需＝缺口 − 已觀測（2024–2025）", "US$B", lambda y: f'=IF({YC[y]}{rows["gap"]}="","",{YC[y]}{rows["gap"]}-{YC[y]}{rows["ext"]})', years=[2024, 2025])
            fr26 = FULL.get((p, 2026), (None,))[0]
            def latest(col):
                a26 = f"Funding!{col}{FDR[(p, '2026')]}" if (p, "2026") in FDR else None
                a25 = f"Funding!{col}{FDR[(p, '2025')]}" if (p, "2025") in FDR else None
                if a26 and a25: return f'IF({a26}<>"",{a26},N({a25}))'
                return f"N({a26 or a25})" if (a26 or a25) else "0"
            pl("lnc", "已簽未起租租賃（最近一期存量；Funding）", "US$B", lambda y: f"={latest(FCOL['lnc'])}", years=[2026], note="G4：只進累計，不分攤到各年")
            pl("jvx", "表外合資承諾（最大損失曝險存量；Funding）", "US$B", lambda y: f"={latest(FCOL['jvx'])}", years=[2026], note="G4：只進累計")
            pl("ext26", "備忘：2026 年已觀測外部融資（Funding 2026 期間，總額口徑；G4 不計入來源）", "US$B",
               lambda y: (f"=N(Funding!{RC['g_ext']}{FDR[(p, '2026')]})" if (p, "2026") in FDR else "=0"), years=[2026])
        pl("cash", "期初現金（In Y9；一次性，只進累計）", "US$B", lambda y: f"=N(gcash_{p}){zs}", years=[2026], note=str(IN.cell(INROW[f"gcash_{p}"], 2).value))
        pl("fcf", "自由現金流＝營業現金流 −（AI＋非 AI）", "US$B", lambda y: f'=IF({YC[y]}{rows["ocf"]}="","",{YC[y]}{rows["ocf"]}-{YC[y]}{rows["need"]})', years=OUTG)
        pl("gr", "缺口占資本支出", "%", lambda y: f'=IF(OR({YC[y]}{rows["gap"]}="",{YC[y]}{rows["need"]}<=0),"",{YC[y]}{rows["gap"]}/{YC[y]}{rows["need"]})', years=OUTG, fmt="0%")
        r += 1

    def verdict(fcf_row, gr_row, miss_expr):
        F = {y: f"{YC[y]}{fcf_row}" for y in range(2026, 2031)}
        R_ = {y: f"N({YC[y]}{gr_row})" for y in range(2026, 2031)}
        consec = "OR(" + ",".join(f"AND({R_[y]}>0.3,{R_[y + 1]}>0.3,{R_[y + 2]}>0.3)" for y in (2026, 2027, 2028)) + ")"
        turn = '"—"'
        for y in (2030, 2029, 2028, 2027, 2026):
            turn = f'IF(AND({",".join(F[z] + ">=0" for z in range(y, 2031))}),"{y}",{turn})'
        return (f'=IF({miss_expr}>0,"未定（資料不足）",IF(OR({F[2030]}<0,{consec}),"OpenAI 型",'
                f'IF(MIN({F[2026]}:{F[2030]})>=0,"Anthropic 型（2026–2030 皆為正）","Anthropic 型（"&{turn}&" 起轉正）")))')
    GPL = [p for p in GP if p != "zhipu"]
    GA.cell(r, 1, "玩家判讀（2026–2030；G3 門檻：2030 自由現金流 < 0，或缺口占資本支出 > 30% 連 3 年＝OpenAI 型；否則 Anthropic 型；營業現金流有缺＝未定）").font = BOLD; r += 1
    for p in GP:
        rows = GPR[p]
        GA.cell(r, 1, f"　{GPN[p]}"); GA.cell(r, 2, f"{p}_verdict")
        GA.cell(r, 4, verdict(rows["fcf"], rows["gr"], f'SUMPRODUCT(({YC[2026]}{rows["ocf"]}:{YC[2030]}{rows["ocf"]}="")*1)'))
        rows["verdict"] = r; r += 1

    # ---- C. 11 家合計 ----
    gsec("C. 主要玩家合計（11 家；智譜不含）")
    SUM11 = lambda k, y: "=" + "+".join(f"N({YC[y]}{GPR[p][k]})" for p in GPL)
    gline("t_ai", "11 家 AI 資料中心資本支出（擁有者口徑）", "US$B", lambda y: SUM11("ai", y), years=OUTG, fill=KEY)
    gline("t_add", "11 家年新增 IT GW（擁有者口徑）", "GW", lambda y: SUM11("add", y), years=OUTG, fmt="0.000")
    gline("t_nai", "11 家非 AI 資本支出", "US$B", lambda y: SUM11("nai", y), years=OUTG)
    gline("t_ocf", "11 家營業現金流（不扣買回股利）", "US$B", lambda y: SUM11("ocf", y), years=OUTG, fill=KEY)
    gline("t_ocfmiss", "　營業現金流無模型值的家數", "家", lambda y: "=" + "+".join(f'({YC[y]}{GPR[p]["ocf"]}="")' for p in GPL), years=OUTG, fmt="0")
    gline("t_ext", "已觀測外部融資（2024–2025）", "US$B", lambda y: "=" + "+".join(f"N({YC[y]}{GPR[p]['ext']})" for p in GPL if "ext" in GPR[p]), years=[2024, 2025])
    gline("t_bb", "備忘：可削減緩衝（買回＋股利）", "US$B", lambda y: SUM11("bb", y), years=OUTG)
    gline("t_cash", "期初現金（11 家合計）", "US$B", lambda y: SUM11("cash", y), years=[2026], fill=KEY,
          note="Microsoft 為 2026-06-30、Oracle 為 2026-08-31（模型最早餘額）；其餘 2025-12-31")
    gline("t_comm", "已簽未起租租賃＋表外合資承諾（存量）", "US$B", lambda y: "=" + "+".join(f"N({YC[y]}{GPR[p]['lnc']})+N({YC[y]}{GPR[p]['jvx']})" for p in GPL if "lnc" in GPR[p]),
          years=[2026], fill=KEY, note="G4：只進累計")
    gline("t_ext26", "備忘：2026 年已觀測外部融資（不計入來源）", "US$B", lambda y: "=" + "+".join(f"N({YC[y]}{GPR[p]['ext26']})" for p in GPL if "ext26" in GPR[p]), years=[2026])
    gline("t_bbcum", "備忘：2026–2030 累計可削減緩衝", "US$B", lambda y: f"=SUM({YC[2026]}{GROW['t_bb']}:{YC[2030]}{GROW['t_bb']})", years=[2026])
    gline("t_ocfmiss2630", "營業現金流無模型值格數（11 家 × 2026–2030；判讀用）", "格", lambda y: f"=SUM({YC[2026]}{GROW['t_ocfmiss']}:{YC[2030]}{GROW['t_ocfmiss']})", years=[2026], fmt="0")

    # ---- D. 整個市場當一間公司（每個需求 × 延遲情景） ----
    gsec("D. 整個市場當一間公司（需求 A／B／C × 延遲 D0／D1／D2；互動頁資料源）。市場其餘＝市場合計 − 11 家（11 家 GW 不隨情景變）；市場其餘沒有資金來源資料（『資金來源未知』）")
    GMK = {}
    for s in SCEN:
        for D in DLY:
            GA.cell(r, 1, f"情景 {s}／{D}").font = BOLD; r += 1
            k0 = f"{s}{D}"
            GMK[k0] = {}
            def ml(k, lab, unit, fy, years=OUTG, fmt="0.0", fill=None):
                GMK[k0][k] = gline(f"m{k0}_{k}", "　" + lab, unit, fy, years=years, fmt=fmt, fill=fill)
            ml("add", "市場年新增 IT GW（Deploy 主路徑，淨）", "GW", lambda y: f"=Deploy!{YC[y]}{DR['na' + s + D]}", years=YEARS, fmt="0.000")
            def mkt(y, k0=k0):
                c = YC[y]; t = [f"{c}{GMK[k0]['add']}*({G('pit', y)}+{GW0}*{G('pfac', y)})"]
                for k_, w_ in ((1, GW1), (2, GW2)):
                    if y + k_ in YC: t.append(f"{w_}*{YC[y + k_]}{GMK[k0]['add']}*{G('pfac', y + k_)}")
                return "=" + "+".join(t)
            ml("mkt", "市場 AI 資料中心資本支出（IT＋廠房提前）", "US$B", mkt)
            ml("own", "需要①：11 家 AI 資料中心資本支出", "US$B", lambda y: f"={G('t_ai', y)}")
            ml("rest", "需要②：市場其餘 AI 資料中心資本支出（資金來源未知）", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['mkt']}-{YC[y]}{GMK[k0]['own']}")
            ml("nai", "需要③：11 家非 AI 資本支出", "US$B", lambda y: f"={G('t_nai', y)}")
            ml("need", "需要合計（①＋②＋③）", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['own']}+{YC[y]}{GMK[k0]['rest']}+{YC[y]}{GMK[k0]['nai']}")
            ml("need12", "需要（不含市場其餘：①＋③）", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['own']}+{YC[y]}{GMK[k0]['nai']}")
            ml("ocf", "來源①：11 家營業現金流（不扣買回股利）", "US$B", lambda y: f"={G('t_ocf', y)}")
            ml("ext", "來源②：已觀測外部融資（只 2024–2025）", "US$B", lambda y: (f"={G('t_ext', y)}" if y <= 2025 else "=0"))
            ml("short", "尚需新融資（不含市場其餘）＝①＋③ − 來源", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['need12']}-{YC[y]}{GMK[k0]['ocf']}-{YC[y]}{GMK[k0]['ext']}", fill=KEY)
            ml("shortR", "尚需新融資（含市場其餘）", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['short']}+{YC[y]}{GMK[k0]['rest']}", fill=KEY)
            c26 = YC[2026]
            ml("cum", "累計尚需新融資（2026 起；不含市場其餘）", "US$B", lambda y, k0=k0: f"=SUM(${c26}{GMK[k0]['short']}:{YC[y]}{GMK[k0]['short']})", years=range(2026, 2031))
            ml("cumnet", "累計仍未安排（不含市場其餘）＝累計 − 期初現金 − 已簽未起租與表外承諾", "US$B",
               lambda y, k0=k0: f"={YC[y]}{GMK[k0]['cum']}-{G('t_cash', 2026)}-{G('t_comm', 2026)}", years=range(2026, 2031), fill=KEY)
            ml("cumcash", "讀法二：累計仍未安排（不含市場其餘；只扣期初現金，承諾不計）", "US$B",
               lambda y, k0=k0: f"={YC[y]}{GMK[k0]['cum']}-{G('t_cash', 2026)}", years=range(2026, 2031))
            ml("cumR", "累計尚需新融資（2026 起；含市場其餘）", "US$B", lambda y, k0=k0: f"=SUM(${c26}{GMK[k0]['shortR']}:{YC[y]}{GMK[k0]['shortR']})", years=range(2026, 2031))
            ml("cumRnet", "累計仍未安排（含市場其餘）", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['cumR']}-{G('t_cash', 2026)}-{G('t_comm', 2026)}", years=range(2026, 2031), fill=KEY)
            ml("cumRcash", "讀法二：累計仍未安排（含市場其餘；只扣期初現金，承諾不計）", "US$B",
               lambda y, k0=k0: f"={YC[y]}{GMK[k0]['cumR']}-{G('t_cash', 2026)}", years=range(2026, 2031))
            ml("fcf12", "自由現金流（只看 11 家）＝營業現金流 −（①＋③）", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['ocf']}-{YC[y]}{GMK[k0]['need12']}")
            ml("gr12", "缺口占資本支出（只看 11 家）", "%", lambda y, k0=k0: f'=IF({YC[y]}{GMK[k0]["need12"]}<=0,"",({YC[y]}{GMK[k0]["need12"]}-{YC[y]}{GMK[k0]["ocf"]})/{YC[y]}{GMK[k0]["need12"]})', fmt="0%")
            ml("fcfR", "自由現金流（含市場其餘）＝營業現金流 − 需要合計", "US$B", lambda y, k0=k0: f"={YC[y]}{GMK[k0]['ocf']}-{YC[y]}{GMK[k0]['need']}")
            ml("grR", "缺口占資本支出（含市場其餘）", "%", lambda y, k0=k0: f'=IF({YC[y]}{GMK[k0]["need"]}<=0,"",({YC[y]}{GMK[k0]["need"]}-{YC[y]}{GMK[k0]["ocf"]})/{YC[y]}{GMK[k0]["need"]})', fmt="0%")
            miss = G("t_ocfmiss2630", 2026)
            for k_, fr_, gr_, lab_ in (("vR", "fcfR", "grR", "判讀（含市場其餘）"), ("v12", "fcf12", "gr12", "判讀（只看 11 家）")):
                GA.cell(r, 1, "　" + lab_); GA.cell(r, 2, f"m{k0}_{k_}")
                GA.cell(r, 4, verdict(GMK[k0][fr_], GMK[k0][gr_], miss)); GMK[k0][k_] = r; GROW[f"m{k0}_{k_}"] = r; r += 1
            r += 1
    GA_LAST = r
    GA.column_dimensions["A"].width = 58; GA.column_dimensions["B"].width = 14; GA.column_dimensions["N"].width = 60
    for y in YEARS: GA.column_dimensions[YC[y]].width = 10

    # ================= Interface：WIF_ =================
    IF["A3"] = "下游取用 2027 年後數值前看覆蓋率（WIF_Coverage_<yyyy>）"; IF["A3"].font = Font(bold=True, color="C00000")
    IF["A1"] = ("Interface — 下游唯一連結點（WIF_ 具名範圍；經 wattonomics/tools/export_snapshot.py 匯成快照）。每格一個名稱：<前綴>_<yyyy>，2024–2032。"
                "單位一律 IT GW（＝TK GW，IT 關鍵電力）；設施電力＝× IF_FacilityGW")
    IF["A2"] = f"版本：瓦帳 {a.version}（{a.date}）；Tokenomics {src['version']}；Siliconomics {ssrc['version']}"
    hdr(IF, 4, ["名稱前綴", "項目", "單位"] + [str(y) for y in OUTY] + ["來源位置"])
    r = 5; GIF = []
    def irow(prefix, lab, sheet, row):
        nonlocal r
        IF.cell(r, 1, prefix); IF.cell(r, 2, lab); IF.cell(r, 3, "GW")
        for j, y in enumerate(OUTY):
            c = IF.cell(r, 4 + j, f"={sheet}!{YC[y]}{row}"); c.font = GREEN; c.number_format = "0.000"
            nm_ = f"{prefix}_{y}"
            name(wb, nm_, f"Interface!${get_column_letter(4 + j)}${r}"); GIF.append(nm_)
        IF.cell(r, 13, f"{sheet}!{YC[2024]}{row}:{YC[2032]}{row}")
        r += 1
    IF.cell(r, 1, "下游取用 2027 年後數值前看覆蓋率（WIF_Coverage_<yyyy>）").font = Font(bold=True, color="C00000"); r += 1
    IF.cell(r, 1, "0. 覆蓋率（有證據支撐的年新增 ÷ 主路徑年新增）").font = BOLD; r += 1
    IF.cell(r, 1, "WIF_Coverage"); IF.cell(r, 2, "覆蓋率（A、D0；< 0.5 時 Checks WARN）"); IF.cell(r, 3, "比")
    for j, y in enumerate(OUTY):
        c = IF.cell(r, 4 + j, f"=Deploy!{YC[y]}{DR['cov']}"); c.font = GREEN; c.number_format = "0.00"
        name(wb, f"WIF_Coverage_{y}", f"Interface!${get_column_letter(4 + j)}${r}"); GIF.append(f"WIF_Coverage_{y}")
    IF.cell(r, 13, f"Deploy!{YC[2024]}{DR['cov']}:{YC[2032]}{DR['cov']}"); IFCOV = r; r += 1
    IF.cell(r, 1, "WIF_Coverage_Power"); IF.cell(r, 2, "覆蓋率（含電力證據路徑；只作參考，Checks 不以此判定）"); IF.cell(r, 3, "比")
    for j, y in enumerate(OUTY):
        c = IF.cell(r, 4 + j, f"=Deploy!{YC[y]}{DR['covp']}"); c.font = GREEN; c.number_format = "0.00"
        name(wb, f"WIF_Coverage_Power_{y}", f"Interface!${get_column_letter(4 + j)}${r}"); GIF.append(f"WIF_Coverage_Power_{y}")
    IF.cell(r, 13, f"Deploy!{YC[2024]}{DR['covp']}:{YC[2032]}{DR['covp']}"); r += 1
    for sfx, nsfx, vlab in (("", "", "主路徑"), ("_SO", "_SiteOnly", "SiteOnly：只有 甲")):
        IF.cell(r, 1, f"A{nsfx}. 市場總 GW（{vlab}；需求情景 × 延遲）").font = BOLD; r += 1
        for s in SCEN:
            for D in DLY:
                irow(f"WIF_MktGW_YE_{s}_{D}{nsfx}", f"年底上線（{s}／{D}；{vlab}）", "Deploy", DR[f"ye{s}{D}{sfx}"])
        for s in SCEN:
            for D in DLY:
                irow(f"WIF_MktGW_Avg_{s}_{D}{nsfx}", f"年平均（{s}／{D}；{vlab}）", "Deploy", DR[f"avg{s}{D}{sfx}"])
        for s in SCEN:
            for D in DLY:
                irow(f"WIF_MktGW_Add_{s}_{D}{nsfx}", f"年新增（淨；{s}／{D}；{vlab}）", "Deploy", DR[f"na{s}{D}{sfx}"])
        IF.cell(r, 1, f"B{nsfx}. 已出貨未上線（{vlab}；相對 D0）").font = BOLD; r += 1
        for s in SCEN:
            for D in ("D1", "D2"):
                irow(f"WIF_ShippedNotOnline_{s}_{D}{nsfx}", f"已出貨未上線（{s}／{D}；{vlab}）", "Deploy", DR[f"sno{s}{D}{sfx}"])
    IF.cell(r, 1, "C. 玩家年底上線 IT GW（主路徑：七家 2027 起＝甲2 累加；營運者口徑；空白＝未知）").font = BOLD; r += 1
    for p in PK:
        if p in XR:
            irow(f"WIF_PlayerGW_YE_{p}", PN[p] + "（主路徑）", "CapexGW", XR[p]["ye"])
        else:
            irow(f"WIF_PlayerGW_YE_{p}", PN[p], "Players", PR[p]["tk"])
    IF.cell(r, 1, "C_SiteOnly. 玩家年底上線 IT GW（只有 甲）").font = BOLD; r += 1
    for p in PK:
        irow(f"WIF_PlayerGW_YE_{p}_SiteOnly", PN[p] + "（SiteOnly）", "Players", PR[p]["tk"])
    def irowU(prefix, lab, sheet, row, unit):
        nonlocal r
        irow(prefix, lab, sheet, row); IF.cell(r - 1, 3, unit)
    IF.cell(r, 1, "D. Spend：年 × 項目市場金額（US$B；主路徑 A、D0 年新增 × 每 GW 內容；上線年口徑）").font = BOLD; r += 1
    for k, rr_ in SPR.items():
        if k == "sil": continue
        irowU(f"WIF_Spend_{k}", str(SP.cell(rr_, 1).value).strip("　"), "Spend", rr_, "US$B")
    IF.cell(r, 1, "E. Facility：每 IT GW 廠房拆分（US$B／IT GW）與市場量").font = BOLD; r += 1
    for k in ("dist", "gen", "ups", "cool", "bldg", "labor", "prof", "net", "land", "resid", "self_unit"):
        irowU(f"WIF_Facility_{k}", str(FC.cell(FCR[k], 1).value).strip("　"), "Facility", FCR[k], "$B/IT GW")
    for k in ("q_gen", "q_self"):
        irowU(f"WIF_Facility_mkt_{k}", str(FC.cell(FMK[k], 1).value).strip("　"), "Facility", FMK[k], "GW")
    IF.cell(r, 1, "F. Funding：資金來源（每家 × 期間；欄位見 Funding 頁；名稱 WIF_Funding_<玩家>_<期間>_<項目>，期間＝2024／2025／2026）").font = BOLD; r += 1
    hdr(IF, r, ["名稱", "玩家", "期間", "值", "項目", "", "", "", "", "", "", "", "來源位置"]); r += 1
    FUNDN = []
    for (pk, b), fr in FDR.items():
        for k, cl in (("capex", RC["den"]), ("int_pct", RC["p_int"]), ("lease_pct", RC["p_fl"]), ("equity_pct", RC["p_eq"]), ("debt_pct", RC["p_debt"]),
                      ("other_pct", RC["p_res"]), ("gross_ext", RC["g_ext"]), ("gross_ext_pct", RC["g_extr"]), ("gross_ocf_pct", RC["g_ocfr"]),
                      ("gross_debt_pct", RC["g_debtr"]), ("gross_equity_pct", RC["g_eqr"]), ("gross_lease_jv_pct", RC["g_flr"]), ("ocf", FCOL["ocf"]), ("fcf", FCOL["fcf"]), ("lnc", FCOL["lnc"]), ("debt_raised", None), ("equity_raised", FCOL["eq"])):
            nm_ = f"WIF_Funding_{pk}_{b}_{k}"
            IF.cell(r, 1, nm_); IF.cell(r, 2, PN[pk]); IF.cell(r, 3, int(b)); IF.cell(r, 5, k)
            if cl is None:
                f = f"=N(Funding!{FCOL['bond']}{fr})+N(Funding!{FCOL['term']}{fr})+N(Funding!{FCOL['gpu']}{fr})+N(Funding!{FCOL['pc']}{fr})"
            else:
                f = f'=IF(Funding!{cl}{fr}="","",Funding!{cl}{fr})'
            c = IF.cell(r, 4, f); c.font = GREEN; c.number_format = "0%" if k.endswith("_pct") else "0.0"
            IF.cell(r, 13, f"Funding!{cl or FCOL['bond']}{fr}")
            name(wb, nm_, f"Interface!$D${r}"); GIF.append(nm_); FUNDN.append(nm_); r += 1
    # ---- G. Gap：資金缺口（v0.1.7；年份 2024–2030；欄 D–J） ----
    IF.cell(r, 1, "G. Gap：資金缺口（US$B；曆年 2024–2030 置於 D–J 欄；名稱 <前綴>_<yyyy>；營業現金流固定為個股模型基準，不隨需求情景變；2028 起覆蓋率 < 0.35 屬情景）").font = BOLD; r += 1
    hdr(IF, r, ["名稱前綴", "項目", "單位"] + [str(y) for y in OUTG] + ["", "", "來源位置"]); r += 1
    GAPN = []

    def grow_if(prefix, lab, row, unit="US$B", years=OUTG, fmt="0.0"):
        nonlocal r
        IF.cell(r, 1, prefix); IF.cell(r, 2, lab); IF.cell(r, 3, unit)
        for y in years:
            j = OUTG.index(y)
            c = IF.cell(r, 4 + j, f'=IF(Gap!{YC[y]}{row}="","",Gap!{YC[y]}{row})'); c.font = GREEN; c.number_format = fmt
            nm_ = f"{prefix}_{y}"; name(wb, nm_, f"Interface!${get_column_letter(4 + j)}${r}"); GIF.append(nm_); GAPN.append(nm_)
        IF.cell(r, 13, f"Gap!{YC[years[0]]}{row}:{YC[years[-1]]}{row}"); r += 1

    def gscalar_if(nm_, lab, ref, unit="US$B", fmt="0.0"):
        nonlocal r
        IF.cell(r, 1, nm_); IF.cell(r, 2, lab); IF.cell(r, 3, unit)
        c = IF.cell(r, 4, f'=IF({ref}="","",{ref})'); c.font = GREEN; c.number_format = fmt
        name(wb, nm_, f"Interface!$D${r}"); GIF.append(nm_); GAPN.append(nm_); IF.cell(r, 13, ref); r += 1
    grow_if("WIF_Gap_PerGW", "每 IT GW 建置成本（IT＋廠房；市場）", GROW["pgw"], "$B/IT GW", fmt="0.00")
    for k, lab, _ in GITEMS + GFAC:
        grow_if(f"WIF_Gap_Item_{k}", f"每 IT GW 項目：{lab}", GROW["i_" + k], "$B/IT GW", fmt="0.00")
    grow_if("WIF_Gap_NonAI", "11 家非 AI 資本支出", GROW["t_nai"])
    grow_if("WIF_Gap_OCF", "11 家營業現金流（不扣買回股利）", GROW["t_ocf"])
    grow_if("WIF_Gap_ExtObs", "已觀測外部融資（Funding 總額口徑）", GROW["t_ext"], years=[2024, 2025])
    grow_if("WIF_Gap_Buyback", "備忘：11 家買回＋股利（可削減緩衝）", GROW["t_bb"])
    grow_if("WIF_Gap_OCFMissing", "11 家營業現金流無模型值家數", GROW["t_ocfmiss"], "家", fmt="0")
    gscalar_if("WIF_Gap_OpenCash", "期初現金（11 家；只進累計）", f"Gap!{YC[2026]}{GROW['t_cash']}")
    gscalar_if("WIF_Gap_Committed", "已簽未起租租賃＋表外合資承諾（存量；只進累計）", f"Gap!{YC[2026]}{GROW['t_comm']}")
    gscalar_if("WIF_Gap_ExtObs2026Memo", "備忘：2026 年已觀測外部融資（不計入來源）", f"Gap!{YC[2026]}{GROW['t_ext26']}")
    gscalar_if("WIF_Gap_BufferCum", "備忘：2026–2030 累計可削減緩衝（買回＋股利）", f"Gap!{YC[2026]}{GROW['t_bbcum']}")
    for s in SCEN:
        for D in DLY:
            k0 = f"{s}{D}"; M = GMK[k0]
            grow_if(f"WIF_Gap_Need_{s}_{D}", f"需要合計（{s}／{D}；11 家 AI＋市場其餘＋11 家非 AI）", M["need"])
            grow_if(f"WIF_Gap_NeedOwn_{s}_{D}", f"需要①：11 家 AI 資料中心（{s}／{D}）", M["own"])
            grow_if(f"WIF_Gap_NeedRest_{s}_{D}", f"需要②：市場其餘（資金來源未知；{s}／{D}）", M["rest"])
            grow_if(f"WIF_Gap_MktAdd_{s}_{D}", f"市場年新增 IT GW（{s}／{D}）", M["add"], "GW", fmt="0.000")
            grow_if(f"WIF_Gap_Short_{s}_{D}", f"尚需新融資（不含市場其餘；{s}／{D}）", M["short"])
            grow_if(f"WIF_Gap_ShortInclRest_{s}_{D}", f"尚需新融資（含市場其餘；{s}／{D}）", M["shortR"])
            grow_if(f"WIF_Gap_CumNeed_{s}_{D}", f"累計尚需新融資（2026 起；不含市場其餘；{s}／{D}）", M["cum"], years=list(range(2026, 2031)))
            grow_if(f"WIF_Gap_CumNeedInclRest_{s}_{D}", f"累計尚需新融資（2026 起；含市場其餘；{s}／{D}）", M["cumR"], years=list(range(2026, 2031)))
            grow_if(f"WIF_Gap_CumShort_{s}_{D}", f"累計仍未安排（不含市場其餘；扣期初現金與承諾；{s}／{D}）", M["cumnet"], years=list(range(2026, 2031)))
            grow_if(f"WIF_Gap_CumShortInclRest_{s}_{D}", f"累計仍未安排（含市場其餘；{s}／{D}）", M["cumRnet"], years=list(range(2026, 2031)))
            grow_if(f"WIF_Gap_CumShortCashOnly_{s}_{D}", f"讀法二：累計仍未安排（不含市場其餘；只扣期初現金；{s}／{D}）", M["cumcash"], years=list(range(2026, 2031)))
            grow_if(f"WIF_Gap_CumShortCashOnlyInclRest_{s}_{D}", f"讀法二：累計仍未安排（含市場其餘；只扣期初現金；{s}／{D}）", M["cumRcash"], years=list(range(2026, 2031)))
            gscalar_if(f"WIF_Gap_Verdict_{s}_{D}", f"判讀（含市場其餘；{s}／{D}）", f"Gap!D{M['vR']}", "文字", fmt="General")
            gscalar_if(f"WIF_Gap_Verdict12_{s}_{D}", f"判讀（只看 11 家；{s}／{D}）", f"Gap!D{M['v12']}", "文字", fmt="General")
    for p in GP:
        rows = GPR[p]
        for k, lab, unit, fmt in (("OCF", "營業現金流", "US$B", "0.0"), ("CapexAI", "AI 資料中心資本支出", "US$B", "0.0"), ("CapexNonAI", "非 AI 資本支出（無模型值以 0 計）", "US$B", "0.0"),
                                  ("Gap", "缺口（正＝需外部資金）", "US$B", "0.0"), ("SelfFund", "自籌比", "x", "0.00"), ("Buyback", "買回＋股利（備忘）", "US$B", "0.0"),
                                  ("CapexAIModel", "對照：個股模型 AI 資本支出", "US$B", "0.0")):
            row = rows[{"OCF": "ocf", "CapexAI": "ai", "CapexNonAI": "nai", "Gap": "gap", "SelfFund": "self", "Buyback": "bb", "CapexAIModel": "cai"}[k]]
            IF.cell(r, 1, f"WIF_Gap_Player_{p}_<yyyy>_{k}"); IF.cell(r, 2, f"{GPN[p]} {lab}"); IF.cell(r, 3, unit)
            for y in OUTG:
                j = OUTG.index(y)
                c = IF.cell(r, 4 + j, f'=IF(Gap!{YC[y]}{row}="","",Gap!{YC[y]}{row})'); c.font = GREEN; c.number_format = fmt
                nm_ = f"WIF_Gap_Player_{p}_{y}_{k}"; name(wb, nm_, f"Interface!${get_column_letter(4 + j)}${r}"); GIF.append(nm_); GAPN.append(nm_)
            IF.cell(r, 13, f"Gap!{YC[2024]}{row}:{YC[2030]}{row}"); r += 1
        gscalar_if(f"WIF_Gap_PlayerCash_{p}", f"{GPN[p]} 期初現金", f"Gap!{YC[2026]}{rows['cash']}")
        if "lnc" in rows:
            gscalar_if(f"WIF_Gap_PlayerCommitted_{p}", f"{GPN[p]} 已簽未起租＋表外合資承諾", f"Gap!{YC[2026]}{rows['lnc']}+Gap!{YC[2026]}{rows['jvx']}")
        gscalar_if(f"WIF_Gap_PlayerVerdict_{p}", f"{GPN[p]} 判讀（2026–2030）", f"Gap!D{rows['verdict']}", "文字", fmt="General")
    IF_LAST = r
    IF.column_dimensions["A"].width = 30; IF.column_dimensions["B"].width = 44; IF.column_dimensions["M"].width = 22

    # ================= Checks =================
    CK["A1"] = "Checks — 恆等式與治理（ERROR 計入 WGOV_Errors，須為 0；WARN 計入 WGOV_Warn，列為待補）"
    hdr(CK, 3, ["ID", "項目", "本庫值", "參照值", "差", "判讀", "性質", "說明"])
    r = 4
    maxr = {"In": IN_LAST + 2, "TK_In": TK_LAST + 2, "SIL_In": 60, "Players": PL_LAST + 2, "Matrix": MX_LAST + 2, "Pipeline": PP_LAST + 2,
            "Deploy": DP_LAST + 2, "C0": C0_LAST + 2, "C1": C1_LAST + 2, "Compare": CP_LAST + 2, "Interface": IF_LAST + 2}
    maxr['CapexGW'] = CX_LAST + 2
    maxr.update({'Funding': FD_LAST + 2, 'FundRaw': FR1 + 2, 'Facility': FC_LAST + 2, 'SegIn': SGSUM + 2, 'Spend': SP_LAST + 2, 'Gap': GA_LAST + 2})
    for sh, mr in maxr.items():
        CK.cell(r, 1, f"E-{sh}"); CK.cell(r, 2, f"{sh} 頁錯誤格數")
        CK.cell(r, 3, f"=SUMPRODUCT(ISERROR({sh}!A1:AT{mr})*1)")
        CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "結構"); r += 1
    for y in (2025, 2026, 2027):
        m = MXR[y]
        CK.cell(r, 1, f"K1-{y}a"); CK.cell(r, 2, f"Matrix {y}：Σ|列和 − Players|")
        CK.cell(r, 3, f"=SUMPRODUCT(ABS(Matrix!{L(SUMC)}{m['r0']}:{L(SUMC)}{m['r1']}-Matrix!{L(PLC)}{m['r0']}:{L(PLC)}{m['r1']}))")
        CK.cell(r, 6, f'=IF(C{r}<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
        CK.cell(r, 1, f"K1-{y}b"); CK.cell(r, 2, f"Matrix {y}：欄和合計（含未分配）＝總量＝Players 市場合計")
        CK.cell(r, 3, f"=SUM(Matrix!C{m['sum']}:{L(UNC)}{m['sum']})"); CK.cell(r, 4, f"=Players!{YC[y]}{TOT['ye']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for y in YEARS:
        CK.cell(r, 1, f"K2-{y}"); CK.cell(r, 2, f"Pipeline D0 {y} 年底＝Players 市場合計（容差 0.01 GW）")
        CK.cell(r, 3, f"=Pipeline!{YC[y]}{PA['yeD0']}"); CK.cell(r, 4, f"=Players!{YC[y]}{TOT['ye']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<=0.01,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for y in OUTY:
        CK.cell(r, 1, f"K3-{y}"); CK.cell(r, 2, f"Pipeline {y} 年底 D0 ≥ D1 ≥ D2")
        CK.cell(r, 3, f"=Pipeline!{YC[y]}{PA['yeD1']}"); CK.cell(r, 4, f"=Pipeline!{YC[y]}{PA['yeD0']}")
        CK.cell(r, 6, f'=IF(AND(Pipeline!{YC[y]}{PA["yeD0"]}>=Pipeline!{YC[y]}{PA["yeD1"]}-1E-9,Pipeline!{YC[y]}{PA["yeD1"]}>=Pipeline!{YC[y]}{PA["yeD2"]}-1E-9),"OK","ERROR")')
        CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K3-2032eq"); CK.cell(r, 2, "延遲只移動時點：D0 與 D2 的 2032 年底差 ≤ 2033 後才上線的部分（2032 均勻型殘差 ×（1 − 2032 年底 D2 比例）＋園區）")
    CK.cell(r, 3, f"=Pipeline!M{PA['yeD0']}-Pipeline!M{PA['yeD2']}")
    CK.cell(r, 4, f"=SUMPRODUCT(ABS(Pipeline!E{AGG['runi']}:M{AGG['runi']}),1-Pipeline!E{FRAC['D2'][2032]}:M{FRAC['D2'][2032]})+SUMPRODUCT(Pipeline!L{CR0}:L{CR1},Pipeline!Q{CR0}:Q{CR1})")
    CK.cell(r, 6, f'=IF(C{r}<=D{r}+1E-9,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    for y in YEARS:
        CK.cell(r, 1, f"K4-{y}"); CK.cell(r, 2, f"{y} 世代占比合計＝1")
        CK.cell(r, 3, f"=Players!{YC[y]}{MIXSUM}"); CK.cell(r, 4, 1)
        CK.cell(r, 6, f'=IF(ABS(C{r}-1)<1E-9,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    for y in OUTY:
        CK.cell(r, 1, f"K5-{y}"); CK.cell(r, 2, f"{y} 殘差＋園區＝Players 年新增")
        CK.cell(r, 3, f"=Pipeline!{YC[y]}{AGG['addchk']}"); CK.cell(r, 4, f"=Players!{YC[y]}{TOT['add']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for k, lo, hi in (("lag_bldg", 0, 24), ("dly_d1_m", 0, 36), ("dly_d1_s", 0, 1), ("dly_d2_m", 0, 36), ("dly_d2_bind", 0, 1), ("tolband", 0, 1),
                      ("k_chip2it", 1, 3), ("g_lt", -0.2, 0.5)):
        for suf in ("", "_lo", "_hi"):
            nm_ = k + suf
            CK.cell(r, 1, f"K6-{nm_}"); CK.cell(r, 2, f"{nm_} 在 {lo}–{hi}")
            CK.cell(r, 3, f"={nm_}"); CK.cell(r, 4, f"{lo}–{hi}")
            CK.cell(r, 6, f'=IF(AND(C{r}>={lo},C{r}<={hi}),"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K7"); CK.cell(r, 2, "未分園區殘差 < 0 的格數（園區合計超過玩家路徑）")
    CK.cell(r, 3, f"=COUNTIF(Pipeline!E{RES[PK[0]]}:M{RES[PK[-1]]},\"<-1E-9\")")
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K8"); CK.cell(r, 2, "IF_FacilityGW 由 Tokenomics 快照提供（0＝快照未含，直接讀活頁簿 SRC_GL_006）")
    CK.cell(r, 3, 1 if HAS_FACGW else 0)
    CK.cell(r, 6, '=IF(C%d=1,"OK","WARN")' % r); CK.cell(r, 7, "治理（TK）")
    CK.cell(r, 8, "Requests R5：請 TK 快照加入 IF_FacilityGW；值相同（1.2，1.1–1.3）"); r += 1
    for y in (2025, 2026, 2027):
        m = MXR[y]
        CK.cell(r, 1, f"K10-{y}"); CK.cell(r, 2, f"Matrix {y}：未分配欄 ≥ 0（分配不超過營運者年底 GW）")
        CK.cell(r, 3, f'=COUNTIF(Matrix!{L(UNC)}{m["r0"]}:{L(UNC)}{m["r1"]},"<-1E-9")')
        CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    for nm_, p, y in (("rng25_neo", "neo", 2025), ("rng25_cn", "cn", 2025), ("rng26_cn", "cn", 2026),
                      ("ovr25_googl", "googl", 2025), ("ovr25_amzn", "amzn", 2025)):
        CK.cell(r, 1, f"K11-{nm_}"); CK.cell(r, 2, f"{nm_} 基準＝Players {p} {y} 年底（區間列與基準一致；設施 ÷ IF_FacilityGW）")
        div = f"/Players!{YC[y]}${FACR}" if p in ("cn", "oth") else ""
        CK.cell(r, 3, f"={nm_}/1000{div}"); CK.cell(r, 4, f"=Players!{YC[y]}{PR[p]['tk']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for y in OUTY:
        CK.cell(r, 1, f"K12-{y}"); CK.cell(r, 2, f"{y} 市場合計 低 ≤ 基準 ≤ 高")
        CK.cell(r, 3, f"=Players!{YC[y]}{TOT['ye']}"); CK.cell(r, 4, f'=TEXT(Players!{YC[y]}{TOT["lo"]},"0.0")&"–"&TEXT(Players!{YC[y]}{TOT["hi"]},"0.0")')
        CK.cell(r, 6, f'=IF(AND(Players!{YC[y]}{TOT["lo"]}<=C{r}+1E-9,C{r}<=Players!{YC[y]}{TOT["hi"]}+1E-9),"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K13"); CK.cell(r, 2, "T2 納入＝1 的園區，引用的證據列在 DB_Evidence 判『錯誤』或『無法核對』的數")
    bad_inc = 0
    evv = {e["id"]: e["verdict"] for e in (json.load(open(EVID, encoding="utf-8")) if os.path.exists(EVID) else [])}
    import re as _re2
    for k, ir in CAMP:
        if str(IN.cell(ir, 14).value) != "1": continue
        ids = _re2.findall(r"[ABD]#\d+", str(IN.cell(ir, 7).value))
        if any(evv.get(i) in ("錯誤", "無法核對") for i in ids[:1]): bad_inc += 1
    CK.cell(r, 3, bad_inc); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理（證據）")
    CK.cell(r, 8, "只檢查說明欄第一個引用（主要依據）"); r += 1
    CK.cell(r, 1, "K14"); CK.cell(r, 2, "DB_Evidence 列數（evidence_raw 全部列）")
    CK.cell(r, 3, len(evv)); CK.cell(r, 6, f'=IF(C{r}>0,"OK","WARN")'); CK.cell(r, 7, "治理（證據）"); r += 1
    CK.cell(r, 1, "K15"); CK.cell(r, 2, "覆蓋率（WIF_Coverage）< 0.5 的年數（2024–2032）")
    CK.cell(r, 3, f'=COUNTIF(Interface!D{IFCOV}:L{IFCOV},"<0.5")'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理（覆蓋）")
    CK.cell(r, 8, "WARN＝該年主路徑新增多半無站點證據（甲2 或收斂規則推得）；下游取用前看覆蓋率"); r += 1
    for y in (2024, 2025, 2026):
        CK.cell(r, 1, f"K16-{y}"); CK.cell(r, 2, f"{y}：主路徑＝SiteOnly（2026 以前只用 甲）")
        CK.cell(r, 3, f"=Deploy!{YC[y]}{DR['yeAD0']}"); CK.cell(r, 4, f"=Deploy!{YC[y]}{DR['yeAD0_SO']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    CK.cell(r, 1, "K17"); CK.cell(r, 2, "甲2 年新增 < 0 的格數（七家 × 年）")
    CK.cell(r, 3, "=" + "+".join(f'COUNTIF(CapexGW!D{XR[p]["gw"]}:M{XR[p]["gw"]},"<0")' for p in XR))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K18"); CK.cell(r, 2, "D1 園區延遲總量（Σ 納入且 asof 後上線的園區 GW ×（t D1 延後 − t D0）× 12，GW·月）> 0（防延遲失效；獨立核對第 1 節）")
    CK.cell(r, 3, f"=SUMPRODUCT(Pipeline!L{CR0}:L{CR1},Pipeline!Q{CR0}:Q{CR1},(Pipeline!K{CR0}:K{CR1}>asof)*1,"
                  f"((Pipeline!J{CR0}:J{CR1}>0)+(Pipeline!J{CR0}:J{CR1}<=0)*dly_d1_s)*(Pipeline!O{CR0}:O{CR1}-Pipeline!M{CR0}:M{CR1})*12)")
    CK.cell(r, 4, f'=SUMPRODUCT(ISTEXT(Pipeline!J{CR0}:J{CR1})*1)')
    CK.cell(r, 6, f'=IF(AND(C{r}>0,D{r}=0),"OK","ERROR")'); CK.cell(r, 7, "驗證")
    CK.cell(r, 8, "D 欄＝J 欄為文字的格數（須為 0）"); r += 1
    for y in (2027, 2029):
        CK.cell(r, 1, f"K19-{y}"); CK.cell(r, 2, f"{y}：SiteOnly D1 < D0（園區＋殘差延遲有作用）")
        CK.cell(r, 3, f"=Deploy!{YC[y]}{DR['yeAD1_SO']}"); CK.cell(r, 4, f"=Deploy!{YC[y]}{DR['yeAD0_SO']}")
        CK.cell(r, 6, f'=IF(C{r}<D{r}-1E-6,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    for y in YEARS:
        CK.cell(r, 1, f"K20-{y}"); CK.cell(r, 2, f"{y} Facility 拆分合計＝TK IF_CapexFacility（世代加權）")
        CK.cell(r, 3, f"=Facility!{YC[y]}{FCR['sum']}"); CK.cell(r, 4, f"=Facility!{YC[y]}{FTK}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    CK.cell(r, 1, "K21"); CK.cell(r, 2, "Facility『其他／未拆』殘差在 −0.5～＋0.5 $B／IT GW（2026；大於表示拆分與 TK 不一致）")
    CK.cell(r, 3, f"=Facility!{YC[2026]}{FCR['resid']}"); CK.cell(r, 6, f'=IF(ABS(C{r})<=0.5,"OK","WARN")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K22"); CK.cell(r, 2, "SegIn 次產業比例合計＝100%（AI半導體 v0.2.3 GB300 對照欄）")
    CK.cell(r, 3, f"=SegIn!C{SGSUM}"); CK.cell(r, 4, 1); CK.cell(r, 6, f'=IF(ABS(C{r}-1)<0.001,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for y in YEARS:
        CK.cell(r, 1, f"K23-{y}"); CK.cell(r, 2, f"{y} Spend IT 設備小計＝年新增 IT GW × 每 IT GW IT capex")
        CK.cell(r, 3, f"=Spend!{YC[y]}{SPR['it_sum']}"); CK.cell(r, 4, f"=Spend!{YC[y]}{SADD}*Spend!{YC[y]}{SIT}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    CK.cell(r, 1, "K24"); CK.cell(r, 2, "Funding：各列內部＋租賃＋股權＋債＋其餘＝分母；其餘 ≥ 0（支付順序）")
    CK.cell(r, 3, "=" + "+".join(f'(ABS(N(Funding!{RC["s_int"]}{fr})+N(Funding!{RC["s_fl"]}{fr})+N(Funding!{RC["s_eq"]}{fr})+N(Funding!{RC["s_debt"]}{fr})+N(Funding!{RC["s_res"]}{fr})-N(Funding!{RC["den"]}{fr}))>1E-9)+(N(Funding!{RC["s_res"]}{fr})<-1E-9)' for fr in FDR.values()))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    n_unver = sum(1 for e in fund["funding"] if e["use"] and e["verdict"] in ("無法核對", "未抽查"))
    CK.cell(r, 1, "K25"); CK.cell(r, 2, "FundRaw 採用＝1 但 G 段判『無法核對』或『未抽查』的列數（建置時）")
    CK.cell(r, 3, n_unver); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理（證據）"); r += 1
    CK.cell(r, 1, "K26"); CK.cell(r, 2, "Anthropic 2026 Funding：洽談中 SPV 36 未進股權／債（洽談中欄＝36、債來源不含）")
    fr = FDR.get(("ant", "2026"))
    CK.cell(r, 3, f"=N(Funding!{FCOL['talks']}{fr})" if fr else 0); CK.cell(r, 4, f"=N(Funding!{FCOL['gpu']}{fr})" if fr else 0)
    CK.cell(r, 6, f'=IF(AND(C{r}=36,D{r}=35),"OK","ERROR")'); CK.cell(r, 7, "驗證（G 段 F191）"); r += 1
    for pk, b, need in (("googl", "2025", 64.564), ("meta", "2025", 30 + 27.3)):
        fr = FDR.get((pk, b))
        CK.cell(r, 1, f"K28-{pk}{b}"); CK.cell(r, 2, f"Funding 總額口徑 {pk} {b} 外部資金 ≥ {need}（GOOGL 發債 F#43；META 發債 F#98＋Hyperion F#101）")
        CK.cell(r, 3, f"=Funding!{RC['g_ext']}{fr}"); CK.cell(r, 4, need)
        CK.cell(r, 6, f'=IF(C{r}>=D{r}-1E-6,"OK","ERROR")'); CK.cell(r, 7, "驗證（核對要求）"); r += 1
    CK.cell(r, 1, "K27"); CK.cell(r, 2, "循環融資表列數＝17（F.md 第三節）")
    CK.cell(r, 3, CIRC1 - CIRC0 + 1); CK.cell(r, 6, f'=IF(C{r}=17,"OK","ERROR")'); CK.cell(r, 7, "結構"); r += 1
    CK.cell(r, 1, "K9"); CK.cell(r, 2, "Deploy：各需求情景 2032 年底 D0 ≥ D1 ≥ D2（情景差額為負時可能不成立）")
    CK.cell(r, 3, "=" + "+".join(f"(Deploy!M{DR['ye' + s + 'D0' + x]}<Deploy!M{DR['ye' + s + 'D1' + x]}-1E-9)+(Deploy!M{DR['ye' + s + 'D1' + x]}<Deploy!M{DR['ye' + s + 'D2' + x]}-1E-9)" for s in SCEN for x in ("", "_SO")))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證"); r += 1
    # ---- Gap（v0.1.7）G-K1–G-K8 ----
    for y in YEARS:
        CK.cell(r, 1, f"G-K1-{y}"); CK.cell(r, 2, f"{y} Gap 每 GW 建置成本項目合計＝每 GW 建置成本")
        CK.cell(r, 3, f"=Gap!{YC[y]}{GROW['isum']}"); CK.cell(r, 4, f"=Gap!{YC[y]}{GROW['pgw']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式（Gap）"); r += 1
    k2 = []
    for p in sorted(XR):
        for y in [y for y in YEARS if y >= 2027]:
            c = YC[y]; g_ = f"CapexGW!{c}{XR[p]['gw']}"
            k2.append(f'IF({g_}="",0,ABS(CapexGW!{c}{XR[p]["fund"]}/(CapexGW!{c}{XR[p]["pit"]}+CapexGW!{c}{XR[p]["pfac"]})-{g_})'
                      f'+ABS(Gap!{c}{GPR[p]["add"]}-{g_})+ABS(Gap!{c}{GPR[p]["pit"]}+Gap!{c}{GPR[p]["pfac"]}-CapexGW!{c}{XR[p]["pit"]}-CapexGW!{c}{XR[p]["pfac"]}))')
    CK.cell(r, 1, "G-K2"); CK.cell(r, 2, "同源一致：CapexGW 當年上線 GW 的資金 ÷ 每 GW 成本＝甲2 年新增 GW；Gap 玩家年新增與每 GW＝CapexGW（七家、2027 起）")
    CK.cell(r, 3, "=" + "+".join(k2)); CK.cell(r, 6, f'=IF(C{r}<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式（Gap）"); r += 1
    M0 = GMK["AD0"]
    CK.cell(r, 1, "G-K3"); CK.cell(r, 2, "市場層 A／D0：11 家擁有＋市場其餘＝市場合計（2024–2030）；11 家擁有＝玩家層加總")
    CK.cell(r, 3, "=" + "+".join(f"ABS(Gap!{YC[y]}{M0['own']}+Gap!{YC[y]}{M0['rest']}-Gap!{YC[y]}{M0['mkt']})+ABS(Gap!{YC[y]}{M0['own']}-(" + "+".join(f"Gap!{YC[y]}{GPR[p]['ai']}" for p in GPL) + "))" for y in OUTG))
    CK.cell(r, 6, f'=IF(C{r}<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式（Gap）"); r += 1
    k4 = []
    for k0, M in GMK.items():
        for y in OUTG:
            c = YC[y]
            k4.append(f"ABS(Gap!{c}{M['short']}-(Gap!{c}{M['own']}+Gap!{c}{M['nai']}-Gap!{c}{M['ocf']}-Gap!{c}{M['ext']}))")
            k4.append(f"ABS(Gap!{c}{M['shortR']}-(Gap!{c}{M['need']}-Gap!{c}{M['ocf']}-Gap!{c}{M['ext']}))")
        for y in range(2026, 2031):
            c = YC[y]; cs = "+".join(f"Gap!{YC[z]}{M['short']}" for z in range(2026, y + 1)); csr = "+".join(f"Gap!{YC[z]}{M['shortR']}" for z in range(2026, y + 1))
            k4.append(f"ABS(Gap!{c}{M['cumnet']}-({cs}-Gap!{YC[2026]}{GROW['t_cash']}-Gap!{YC[2026]}{GROW['t_comm']}))")
            k4.append(f"ABS(Gap!{c}{M['cumRnet']}-({csr}-Gap!{YC[2026]}{GROW['t_cash']}-Gap!{YC[2026]}{GROW['t_comm']}))")
    CK.cell(r, 1, "G-K4"); CK.cell(r, 2, "每情景：尚需新融資＝需要 − 來源；累計仍未安排＝逐年加總 − 期初現金 − 承諾（9 情景 × 年）")
    CK.cell(r, 3, "=" + "+".join(k4)); CK.cell(r, 6, f'=IF(C{r}<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式（Gap）"); r += 1
    import re as _re3
    nbad5 = 0
    for p, (rg, rsf, rbb, rbbf) in GAPREF.items():
        for y in OUTG:
            for rr_ in (rg, rsf, GPR[p]["need"]):
                f = str(GA.cell(rr_, YEARS.index(y) + 4).value)
                refs = {int(x) for x in _re3.findall(r"(?<![A-Za-z!])[A-M]\$?(\d+)", f)}
                if refs & {rbb, rbbf}: nbad5 += 1
    for k0, M in GMK.items():
        for y in OUTG:
            f = str(GA.cell(M["short"], YEARS.index(y) + 4).value) + str(GA.cell(M["shortR"], YEARS.index(y) + 4).value)
            if str(GROW["t_bb"]) in {x for x in _re3.findall(r"(?<![A-Za-z!])[A-M]\$?(\d+)", f)}: nbad5 += 1
    CK.cell(r, 1, "G-K5"); CK.cell(r, 2, "缺口、自籌比、尚需新融資的公式不引用買回＋股利列（建置時公式文字檢查；Andy 2026-10-31 04:20：不扣買回與股利）")
    CK.cell(r, 3, nbad5); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "規格（Gap）"); r += 1
    CK.cell(r, 1, "G-K6a"); CK.cell(r, 2, "市場其餘 AI 資本支出 < 0 的情景 × 年數（11 家 > 市場；2024–2030）")
    CK.cell(r, 3, "=" + "+".join(f'COUNTIF(Gap!{YC[2024]}{M["rest"]}:{YC[2030]}{M["rest"]},"<-1E-9")' for M in GMK.values()))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證（Gap）"); r += 1
    CK.cell(r, 1, "G-K6b"); CK.cell(r, 2, "11 家營業現金流『無模型值』的公司 × 年數（2024–2030）")
    CK.cell(r, 3, f"=SUM(Gap!{YC[2024]}{GROW['t_ocfmiss']}:{YC[2030]}{GROW['t_ocfmiss']})")
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證（Gap）"); r += 1
    k7 = []
    for p in sorted(XR):
        for y in (2024, 2025, 2026):
            c = YC[y]; x_ = f"Gap!{c}{GPR[p]['capx']}"; a_ = f"Gap!{c}{GPR[p]['ai']}"
            k7.append(f'IF(OR({x_}="",N({x_})=0),0,ABS({a_}/{x_}-1)>tolband)')
    CK.cell(r, 1, "G-K7"); CK.cell(r, 2, "Gap 七家 AI 資本支出（新增 GW × 每 GW）÷ CapexGW 曆年 AI 資本支出（In Y2 × AI 占比）超出 tolband 的家 × 年數（2024–2026；只警告）")
    CK.cell(r, 3, "=" + "+".join(f"({x})" for x in k7)); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證（Gap）")
    CK.cell(r, 8, "2024–2025 年新增多為『新納入存量』（不計資本支出），推算值偏低；見報告"); r += 1
    CK.cell(r, 1, "G-K8"); CK.cell(r, 2, "市場層 A／D0 IT 部分（年新增 × 每 GW IT）＝Spend IT 設備小計（2024–2032；同一組儲存格）")
    CK.cell(r, 3, "=" + "+".join(f"ABS(MAX(0,Gap!{YC[y]}{M0['add']})*Gap!{YC[y]}{GROW['pit']}-Spend!{YC[y]}{SPR['it_sum']})" for y in OUTY))
    CK.cell(r, 6, f'=IF(C{r}<1E-6,"OK","ERROR")'); CK.cell(r, 7, "恆等式（Gap）"); r += 1
    a0, a1 = 5, IN_LAST
    CK.cell(r, 1, "G1"); CK.cell(r, 2, "In 頁輸入列缺標記")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}<>"")*(In!E{a0}:E{a1}=""))'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G2"); CK.cell(r, 2, "In 頁輸入列缺來源 ID")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}<>"")*(In!F{a0}:F{a1}=""))'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G3"); CK.cell(r, 2, "純量列標記含 Assumed／Analogy 但未給低／高")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}=1)*((ISNUMBER(SEARCH("Assumed",In!E{a0}:E{a1}))+ISNUMBER(SEARCH("Analogy",In!E{a0}:E{a1})))>0)*(((In!H{a0}:H{a1}="")+(In!J{a0}:J{a1}=""))>0))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G3b"); CK.cell(r, 2, "年度／分攤列標記含 Assumed 或 Analogy（年度列無低／高欄，v0.2 補區間列）")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}>1)*((ISNUMBER(SEARCH("Assumed",In!E{a0}:E{a1}))+ISNUMBER(SEARCH("Analogy",In!E{a0}:E{a1})))>0))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理"); r += 1
    keys = list(INROW)
    CK.cell(r, 1, "G4"); CK.cell(r, 2, "In 頁鍵重複（建置時檢查）"); CK.cell(r, 3, len(keys) - len(set(keys)))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G5"); CK.cell(r, 2, "標『待證據』的輸入列數")
    CK.cell(r, 3, f'=COUNTIF(In!E{a0}:E{a1},"待證據")'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理")
    CK.cell(r, 8, "evidence_raw 複核後入 In，逐步歸零"); r += 1
    CK.cell(r, 1, "G6"); CK.cell(r, 2, "延遲參數仍為佔位值（說明含【佔位】）")
    CK.cell(r, 3, f'=COUNTIF(In!G{a0}:G{a1},"*【佔位】*")'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理"); r += 1
    r += 1
    CK.cell(r, 1, "WGOV_Errors").font = BOLD; c = CK.cell(r, 3, f'=COUNTIF(F4:F{r - 2},"ERROR")'); c.fill = KEY
    name(wb, "WGOV_Errors", f"Checks!$C${r}"); r += 1
    CK.cell(r, 1, "WGOV_Warn").font = BOLD; c = CK.cell(r, 3, f'=COUNTIF(F4:F{r - 3},"WARN")'); c.fill = KEY
    name(wb, "WGOV_Warn", f"Checks!$C${r}")
    for col, w in zip("ABCDEFGH", (12, 70, 12, 12, 10, 8, 12, 60)):
        CK.column_dimensions[col].width = w

    # ================= Decisions、DB_Evidence、Sources、Requests =================
    DC["A1"] = "Decisions — 決策登錄（GL01-01 起）。判斷類由 Andy 決定；工程細節依預設執行者標『預設、可撤回』"
    hdr(DC, 3, ["決策 ID", "類別", "項目", "決議", "日期", "依據", "狀態"])
    decisions = list(seed["decisions"]) + [
        {"id": "GL01-11", "cat": "工程", "item": "Rubin Ultra 占比", "decision": "Units 顆數 ÷ TK 每 GW GPU 數換成 GW 占比（RU ÷（VR200＋RU））；2027 起 Blackwell 不計", "date": "2026-10-28", "quote": "工作單第 2 節", "status": "預設、可撤回"},
        {"id": "GL01-12", "cat": "待決", "item": "延遲參數 D1／D2、g_lt、lag_bldg、2023–25 NonNV 比例", "decision": "佔位值（In S1、S2；Y6 nnv_mkt）；待 B_delay.csv 與 Andy 判斷", "date": "2026-10-28", "quote": "chat 端建置指示", "status": "預設、可撤回"},
        # v0.1.7 Gap 頁（工作單 20261031 §1 判斷題 G1–G6 與建置判斷）
        {"id": "GL01-25", "cat": "判斷", "item": "G1 成本口徑", "decision": "主口徑＝資本支出（現金；新增 GW × 每 GW，廠房提前 lag_bldg）；年化持有成本只作對照列——TK 快照未含持有成本名稱，本版列『待 TK』（Requests R6），不自行估算", "date": "2026-10-31", "quote": "規劃書 §5 Q1；Andy 10:22「每座 data center 要花多少錢」", "status": "預設、可撤回"},
        {"id": "GL01-26", "cat": "判斷", "item": "G2 非 AI 資本支出", "decision": "計入『需要』但獨立成一層（11 家非 AI 資本支出）；來源＝個股模型非 AI 列；模型沒拆或該年沒有（Alphabet、Meta、Oracle FY2026，CoreWeave、Nebius、WhiteFiber、OpenAI、Anthropic 全期）標無模型值、以 0 計", "date": "2026-10-31", "quote": "工作單 §1 G2", "status": "預設、可撤回"},
        {"id": "GL01-27", "cat": "判斷", "item": "G3 判讀門檻", "decision": "2026–2030：2030 年自由現金流 < 0，或缺口占資本支出 > 30% 連 3 年＝OpenAI 型；否則 Anthropic 型（標轉正年）；營業現金流有缺＝未定（資料不足）。市場層出兩讀：含市場其餘／只看 11 家", "date": "2026-10-31", "quote": "規劃書 §5 Q4；工作單 §1 G3", "status": "預設、可撤回"},
        {"id": "GL01-28", "cat": "判斷", "item": "G4 前瞻年份已到位資金", "decision": "2024–2025 用 Funding 已觀測外部融資（總額口徑；Microsoft、Oracle 財年按月數換曆年）；2026 起年度來源只有營業現金流；期初現金（模型 2025 年底，Microsoft 2026-06、Oracle 2026-08 為模型最早餘額）與尚未起租租賃＋表外合資最大損失曝險（Funding 最近一期存量）只進累計。2026 年已觀測融資只列備忘", "date": "2026-10-31", "quote": "工作單 §1 G4", "status": "預設、可撤回"},
        {"id": "GL01-29", "cat": "判斷", "item": "G5 營業現金流不隨需求情景", "decision": "是；來源固定為個股模型基準情境（B）", "date": "2026-10-31", "quote": "工作單 §1 G5", "status": "預設、可撤回"},
        {"id": "GL01-30", "cat": "判斷", "item": "G6 年份", "decision": "2024–2030；2028 起標『情景（覆蓋率 < 0.35）』；個股模型沒有的年份標無模型值、不外推；2024–2025 營業現金流模型沒有時改用 Funding 全年觀測（期間標籤 CY／FY），仍沒有則無模型值", "date": "2026-10-31", "quote": "規劃書 §2、§4、§6", "status": "預設、可撤回"},
        {"id": "GL01-31", "cat": "工程", "item": "OpenAI／Stargate 自建 GW 的擁有者", "decision": "Players 的 oai（Stargate 自建，SoftBank、SB Energy 營運）依規劃書 §3.3 視為 SPV 擁有，歸『市場其餘』，不計 OpenAI 資本支出；OpenAI 營業現金流＝模型自由現金流＋模型自有資本支出（算力現金成本內含）", "date": "2026-10-31", "quote": "規劃書 §3.3（表外合資與 SPV）", "status": "預設、可撤回"},
        {"id": "GL01-32", "cat": "工程", "item": "無 CFO 列的模型", "decision": "CoreWeave、Nebius、WhiteFiber：營業現金流＝營運缺口＋CapEx＋JV 出資＋overlay（WhiteFiber 另加回股利）＝營運來源 − 租金 − 利息（− 現金稅）；Anthropic：F07＋F03；OpenAI：F06＋Cost 自有資本支出；智譜：F09（RMB 億 ÷ 10 ÷ gfx_cny）", "date": "2026-10-31", "quote": "工作單 §2e（逐家人工確認列位）", "status": "預設、可撤回"},
        {"id": "GL01-33", "cat": "工程", "item": "Gap 每 GW 成本", "decision": "市場：IT＝Spend 頁同式（市場 NonNV × spend_nnv_hold）；廠房＝CapexGW 同式套市場 NonNV（NVIDIA 世代加權 IF_CapexFacility 與 VR200 × spend_nnv_hold），項目依 Facility 拆分比例。玩家：CapexGW（七家）或 C1 同式（其餘）。市場其餘＝市場合計 − 11 家（含每 GW 成本差）", "date": "2026-10-31", "quote": "工作單 §2a", "status": "預設、可撤回"},
    ]
    r = 4
    for d in decisions:
        for j, k in enumerate(("id", "cat", "item", "decision", "date", "quote", "status"), 1):
            c = DC.cell(r, j, d.get(k)); c.alignment = WRAP
        if "預設" in str(d.get("status")):
            for j in range(1, 8): DC.cell(r, j).fill = PEND
        r += 1
    for col, w in zip("ABCDEFG", (9, 8, 24, 80, 12, 34, 14)):
        DC.column_dimensions[col].width = w
    EV["A1"] = "DB_Evidence — 證據登錄（wattonomics/evidence_raw/*.csv 經 chat 端複核後整理入本頁；含不採納者）"
    EV["A2"] = ("由 tools/import_evidence.py 自 evidence_raw 產生；C 段判定以 C_review.md 為準。只有『屬實』與『小誤（已更正）』可入 In；"
                "『錯誤』『無法核對』『未抽查』只列於此（T2 納入＝0 或只作說明）。對應鍵＝In 頁說明欄引用本列的鍵")
    hdr(EV, 3, ["ID", "玩家／範圍", "主張（摘要）", "值", "口徑", "as_of", "標記（原檔）", "C 段判定", "C 段說明／更正", "對應 In 鍵", "採用", "網址"])
    import re as _re
    notes = {k: str(IN.cell(rr, 7).value or "") + " " + str(IN.cell(rr, 2).value or "") for k, rr in INROW.items()}
    ev = json.load(open(EVID, encoding="utf-8")) if os.path.exists(EVID) else []
    r = 4
    for e in ev:
        pat = _re.compile(_re.escape(e["id"]) + r"(?!\d)")
        used = [k for k, nt in notes.items() if pat.search(nt)]
        inc = [k for k in used if not (INKIND.get(k) == 4 and str(IN.cell(INROW[k], 14).value) != "1")]
        vals = (e["id"], e["player"], e["claim"], e["value"], e["basis"], e["as_of"], e["tag"], e["verdict"], e["cnote"], "、".join(used),
                ("採用" if inc else "引用（未入值）") if used else "—", e["url"])
        for j, v in enumerate(vals, 1):
            EV.cell(r, j, v)
        if e["verdict"] in ("錯誤", "無法核對"):
            for j in range(1, 13): EV.cell(r, j).fill = PEND
        if inc and e["verdict"] in ("錯誤", "無法核對", "未抽查"):
            EV.cell(r, 11, "引用（見 In 說明：值另有 repo 來源，或為更正／排除註記）")
        r += 1
    EV_LAST = r - 1
    for col, w in zip("ABCDEFGHIJKL", (7, 18, 50, 16, 9, 11, 22, 12, 40, 24, 14, 40)):
        EV.column_dimensions[col].width = w
    SR["A1"] = "Sources — 來源 ID（In 頁 F 欄）"
    hdr(SR, 3, ["來源 ID", "路徑", "說明"])
    r = 4
    for s in seed["sources"]:
        SR.cell(r, 1, s["id"]); SR.cell(r, 2, s["path"]); SR.cell(r, 3, s["desc"]); r += 1
    for p in GP:   # v0.1.7 Gap 頁個股模型（In Y9）
        g = GSEED["players"].get(p, {})
        SR.cell(r, 1, GSRC[p]); SR.cell(r, 2, g.get("file", ""))
        SR.cell(r, 3, f"{GPN[p]} 個股模型（main commit {g.get('commit', '')}；{g.get('family', '')}；基準情境）— Gap 頁 In Y9 取數（tools/import_gap_models.py，逐家人工確認列位）"); r += 1
    for col, w in zip("ABC", (14, 70, 80)): SR.column_dimensions[col].width = w
    RQ["A1"] = "Requests — 需要 TK 或 Siliconomics 權威改值或新增名稱的請求（本庫不改其檔案；工作單第 4 節停止條件）"
    hdr(RQ, 3, ["ID", "對象", "請求", "原因", "狀態"])
    reqs = [("R1", "Tokenomics", "確認每 GW 口徑", "已裁定（2026-10-28）：TK Interface 每 GW＝IT 關鍵電力；IF_FacilityGW 第 8 列＝每 IT GW 對應設施電力 1.2（1.1–1.3）；瓦帳主口徑改 IT GW、g_ITtoFac＝1", "已裁定"),
            ("R2", "Siliconomics", "Deploy 頁標籤", "撤回（規劃書 r1：IT 電力標籤正確）", "撤回"),
            ("R3", "Tokenomics", "NonNV 每 GW 顆數與每顆功率的 IF_ 名稱（v0.2 C2）", "規劃書第 5 節", "v0.2"),
            ("R4", "瓦帳 v0.2（C2）", "晶片合約對帳：OpenAI–AMD 6 GW（晶片功耗）、Broadcom 10 GW、NVIDIA Vera Rubin 5 GW；Anthropic–AMD 2 GW、Google／Broadcom TPU 3.5（Anthropic 稱 5）GW", "B_summary 第 0 節第 2 點：晶片合約與機房合約是同一批容量的兩面，不入 C0", "v0.2"),
            ("R5", "Tokenomics", "快照（tokenomics_snapshot）加入 IF_FacilityGW", "v21.7 快照未含；瓦帳暫直接讀活頁簿 Interface!C8:Q8（commit 730b6a5）", "待 TK"),
            ("R6", "Tokenomics", "快照加入年化持有成本（AL_HoldGen 或 IF_HoldGen，世代別）", "v0.1.7 Gap 頁 G1 對照列需要（年化持有成本 × 年平均 GW）；v21.7 快照未含，本版不列數值", "待 TK")]
    r = 4
    for q in reqs:
        for j, v in enumerate(q, 1): RQ.cell(r, j, v).alignment = WRAP
        r += 1
    for col, w in zip("ABCDE", (6, 14, 60, 70, 14)): RQ.column_dimensions[col].width = w

    # ================= README =================
    lines = [
        (f"瓦帳 {a.version} — 全市場 AI 資料中心 GW 唯一基準事實（{a.date}）", True),
        ("命題：市場總 GW 要與主要玩家 GW 對得上，再與資本支出、GPU／ASIC 顆數、記憶體、電廠與發電機組、電源散熱、AI 投資地圖個股逐環對得上；並加入與需求情景正交的基建延遲情景。", False),
        ("驅動 → 推導：In（Epoch 晶片功耗、玩家年底 MW、觀測、園區管線、延遲 D0–D2、世代組合、使用者分配）＋ TK_In（每 GW 係數、IF_FacilityGW）＋ SIL_In（顆數、情景係數）→ Players（玩家 × 年 IT GW）→ Matrix（營運者 × 使用者）→ Pipeline（上線事件 × 延遲）→ Deploy（需求 × 延遲 3×3）→ C0／C1／Compare（對帳）→ Interface（WIF_）→ Checks（WGOV_Errors＝0）。", False),
        ("單位（規劃書 r1）：主口徑 IT GW＝TK GW（IT 關鍵電力）；設施電力＝IT × TK IF_FacilityGW（1.2）。公司標設施者 ÷ IF_FacilityGW；口徑未說明者基準取 IT 讀法，Players 另列設施讀法；Epoch 晶片功耗 × k_chip2it。", False),
        ("頁面：README｜In 輸入｜TK_In、SIL_In 鏡像（綠字）｜Players 玩家 × 年｜Matrix 營運者 × 使用者｜Pipeline 園區管線 × 延遲｜Deploy 市場總 GW 需求 × 延遲｜C0 由下而上 vs 第三方｜C1 資本支出對帳｜Compare 與 Siliconomics 成本等值 GW｜Funding、FundRaw 資金來源｜Gap 資金缺口（v0.1.7）｜Facility、SegIn、Spend｜Interface WIF_｜Checks｜Decisions｜DB_Evidence｜Sources｜Requests。", False),
        ("v0.1 範圍：C0、C1、延遲管線（D0–D2）、WIF_ 輸出；C2–C7 自 v0.2 起。第 2 階段：填入 evidence_raw 證據（只用 C 段判屬實或已更正的列；DB_Evidence 頁列全部證據與判定），延遲參數依 B_delay 校準。", False),
        ("第 3 階段：新增 CapexGW 頁（甲2：資本支出推得 GW；MSFT、GOOGL、AMZN、META、ORCL、CRWV、NBIS 2027 起進主路徑）；Interface 另出 _SiteOnly（只有 甲）；WIF_Coverage_<yyyy>。下游取用 2027 年後數值前看覆蓋率。", True),
        ("v0.1.1（改名 Wattonomics／瓦帳）：新增 Funding（玩家 × 年資金來源、循環融資 17 條）、FundRaw（F 段證據＋G 段判定）、Facility（每 IT GW 廠房拆分與市場層、自建發電）、SegIn（AI半導體 次產業比例鏡像）、Spend（年 × 項目市場金額）；C1 加廣義口徑（資本支出＋尚未起租變化）；Interface 增 WIF_Funding_*、WIF_Spend_*、WIF_Facility_*。", True),
        ("v0.1.2：套用 H 段（API 旗標人工複核）更正——證據檔 79 筆、In 輸入 19 格（cp_msft_abl 改 336＋336 MW-IT、xfmr_gva 2.1／2.67／3.15、fac_gen_px 低 0.5），FundRaw F#31、F#199、F#222 改不採用（採購或投資存量，非融資）；"
         "園區候選（S 段 409 列，site_pipeline）經 API 回查與 I 段未參與代理開原文抽查後合併：新增 17 列（計入 13 列）、更正既有 20 列（Fairwater 拆分、Ports Pike 首批 800 MW 至 2031、Humain 1.7 GW 提前到 2027、Milam 308＋445、Jupiter 已知延遲刪除）、3 列改計入（Stargate UAE、Colossus 2、Mistral 96 MW）。Hyperion 5 GW、SoftBank 苫小牧依判斷題預設不入。", True),
        ("v0.1.3：第二輪美國以外園區候選（T 段 345 列，site_pipeline 分區 Batch）經 API 回查與 I2 段未參與代理開原文抽查後合併：新增 10 列（計入 5：Nebius Beit Shemesh、Google Kronstorf、DataVita DV1 擴建、Meta Sturgeon 250＋720 設施）；Humain 兩列加註風險（只有公司目標支持）；主權玩家名稱更新。", True),
        ("v0.1.4：Q3 財報 F#255–257（AWS backlog、Amazon 2027 資本支出未量化指引、Meta 9-30 債務與現金）。", True),
        ("v0.1.5：補 2027 覆蓋率（J 段）：Epoch 站點時間線逐頁、公司與電力公司文件、AI半導體 2026-10-30 送件，經未參與代理開原文複核（30 列可用、1 列無法核對）後合併——新增 T2 27 列（Google Fort Wayne、Cedar Rapids、Storey、Columbus、Lincoln、Goodnight；AWS New Albany、Ridgeland B5、Berwick、Black Pearl DH4；Meta Montgomery、Cheyenne、Bowling Green；Oracle Saline 第 1 棟、Red Oak；Nebius Mäntsälä；Nscale Sines、Lake Mariner；2028：Van Buren、Kozani），改既有 6 列（Shackelford 依 Epoch 時程前移、Saline 1,383→922、Lordstown 納入、Nebius PA 260／Lappeenranta 70 納入）。WIF_Coverage_2027 0.377→0.526；只動 2027 年以後的上線量，2026 年底不變。", True),
        ("v0.1.6：K 段 2028H2–2032 證據（未參與代理複核：屬實 26、小誤 9）——新增 T2 33 列：站點 10 列（Microsoft Pecos 2,000、Catawba 720 IT、Rome 422、Aragón 150；Google Visakhapatnam 1,000；CoreWeave Helios III 132）、電力證據路徑 10 列（納入＝2：NIPSCO–Amazon 確定 1,200、Georgia Power–OpenAI 確定 2,200；只進 WIF_Coverage_Power，不入 Players）、未核准合約 13 列（納入＝0）。新名稱 WIF_Coverage_Power_<yyyy>。2026 年以前 0 變動，2027 年主路徑與 SiteOnly 不變。", True),
        ("v0.1.7：新增 Gap 頁（資金缺口；Andy 2026-10-31 規格）：需要的錢＝新增 IT GW × 每 IT GW 建置成本（TK IF_CapexTotal；IT 與 Spend 同式、廠房與 CapexGW 同式；廠房提前 lag_bldg），每 GW 只在擁有者計一次；"
         "錢從哪來＝11 家全公司營業現金流（個股模型基準情境，不扣買回股利）＋已觀測外部融資（2024–2025）；『整個市場當一間公司』逐情景列需要三層（11 家 AI、市場其餘、11 家非 AI）、來源兩層、尚需新融資與累計仍未安排；In 新節 Y9（個股模型基準值）；"
         "Interface 新增 WIF_Gap_*；既有 WIF_ 名稱 0 變動。判斷題 G1–G6 見 Decisions GL01-25 至 GL01-33。", True),
        ("顏色：藍字＝輸入；黑字＝公式；綠字＝跨頁或外部鏡像；淡黃底＝關鍵輸出；亮黃底＝待證據或預設待決。", False),
        ("Excel 優先：要改輸入，直接改 In 頁藍字；builder 重建時讀回（--base）。公式頁不要手改。下游只讀 Interface 的 WIF_ 名稱（經快照）。", False),
        ("版本命名：YYYYMMDD_Wattonomics_vX.Y.xlsx；model/CURRENT 記錄現行檔名；data/CURRENT_SNAPSHOT 指向現行快照。", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = RD.cell(i, 1, t); c.alignment = WRAP
        if b: c.font = Font(bold=True, size=13)
    RD.column_dimensions["A"].width = 150

    wb.save(a.out)
    print(f"寫出 {a.out}：頁面 {len(wb.sheetnames)}、In 鍵 {len(INROW)}、GIF 名稱 {len(GIF)}、具名範圍 {len(wb.defined_names)}")


if __name__ == "__main__":
    main()
