# 奇鋐（3017）收支與估值模型 v0.2（AI 半導體 Project：電源與散熱次產業，樣板＝delta/）
# 用法：python3 avc/build_avc.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；市場與每架內容以美元輸入、以匯率 fx 換算；EPS：新台幣。
# 口徑：歸屬母公司淨利（合併淨利 − 非控制權益；富世達為子公司）。
import sys, os, glob
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
DEL_F = only("delta/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ITC = None
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    if r[0].value == "AI IT 資本支出（加速器 × 每顆 IT capex）":
        ITC = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert ITC and all(v is not None for v in ITC[0]), ITC
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("racks", "capit", "cool"):
        BR[k] = (r[4].value, r[7].value, f"{rel(BR_F)}｜In!E{r[0].row}（GB300）、H{r[0].row}（VR200）")
assert len(BR) == 3, BR.keys()
DLC = None
for r in openpyxl.load_workbook(DEL_F, data_only=True)["Units"].iter_rows(min_row=4):
    if r[0].value == "台達電液冷營收":
        DLC = ([r[2 + j].value for j in range(4)], f"{rel(DEL_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert DLC and all(v is not None for v in DLC[0]), DLC

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B），市場與每架內容為美元；比例：%。每個輸入盡量附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝奇鋐法說（經鉅亨、優分析、時報轉述）與月營收公告；公司的產品比重說法為 Interested-party；券商估計皆為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
UA4Q = "https://uanalyze.com.tw/articles/1131748834"
UA2Q = "https://uanalyze.com.tw/articles/5815653290"
WR1Q = "https://wantrich.chinatimes.com/news/20260514900569-420101"
WR2Q = "https://wantrich.chinatimes.com/news/20260813900226-420101"
KGI08 = "https://www.kgi.com.hk/zh-cn/-/media/files/kgishk/research-reports/tw-reports/asia-vital-components-3017-tt20260812en.pdf"
KGI05 = "https://www.kgi.com.hk/zh-cn/-/media/files/kgishk/research-reports/tw-reports/asia-vital-components_20260805_en.pdf"
TN09 = "https://finance.technews.tw/2026/10/07/avc-3017-202609-financial-report/"
TN1019 = "https://finance.technews.tw/2026/10/19/avc-target-price/"
TN0429 = "https://finance.technews.tw/2026/04/29/avc-ai-server/"
ETTF = "https://finance.ettoday.net/news/3220835"
HIS = "https://histock.tw/stock/3017/%E5%88%A9%E6%BD%A4%E6%AF%94%E7%8E%87"

sec("A. AI 部署規模（讀 markets/network、markets/chip_bridge、tsmc 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, ITC[0][j], None, "Derived", ITC[1],
        "加速器顆數 × 每顆 IT capex（含 ASIC）；本模型以此除以每 GW IT 資本支出得到部署 GW" if y == 26 else "同上")
inp("itpg_gb", "每 GW IT 資本支出（GB300）", "US$B", None, BR["capit"][0], None, "Derived", BR["capit"][2], "Tokenomics v20.1 IF_CapexIT（經橋接表；唯讀）")
inp("itpg_vr", "每 GW IT 資本支出（VR200；2027 起）", "US$B", None, BR["capit"][1], None, "Derived", BR["capit"][2], "同上；Rubin Ultra 與 ASIC 機架未另訂（Assumed：同 VR200）")
inp("vr26", "2026 年部署中 Rubin 世代占比（混合每 GW 數值用）", "%", *SIL.vals("vr26"), "Interested-party", "TrendForce 2026 組合 Rubin 22%（經 markets/memory；與 delta 模型相同）", "公司 2Q26 法說：GPU 量在 3Q 末放大，Vera Rubin 下半年推出（時報 2026-06-02 股東會）")
inp("racks_gb", "每 GW 機架數（GB300）", "架", None, BR["racks"][0], None, "Derived", BR["racks"][2], "Tokenomics（經橋接表）")
inp("racks_vr", "每 GW 機架數（VR200）", "架", None, BR["racks"][1], None, "Derived", BR["racks"][2], "同上")
inp("cool_gb", "GB300 每架機架內液冷（運算匣＋交換匣，橋接表）", "US$K", None, BR["cool"][0], None, "Interested-party", BR["cool"][2], "Morgan Stanley（經 Tom's Hardware 2025-11-05）：GB300 NVL72 冷卻 BOM US$49.9K；不含 CDU、分歧管")
inp("cool_vr", "VR200 每架機架內液冷（橋接表）", "US$K", None, BR["cool"][1], None, "Interested-party", BR["cool"][2], "同上：VR NVL144 US$55.7K（運算匣每匣 2,660、交換匣每匣 870）")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、台達電模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
inp("fx", "匯率（新台幣 ÷ 美元）", "NT$/US$", 30.5, 32.0, 33.0, "Derived", "與 delta 模型相同（公司 2Q26 法說換算隱含約 32.7）", "凱基 2026-08-12：市值 NT$1,148bn ≈ US$33.6bn（隱含約 34.2，口徑待查）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"dlc_{y}", f"台達電液冷營收 20{y}（台達電模型，情境 A；Recon 用）", "US$B", None, DLC[0][j], None, "Derived", DLC[1],
        "只作對帳（台達電液冷多在機架外：CDU、Liquid-to-Air），不進本模型計算" if y == 26 else "同上")

sec("B. 液冷內容、滲透率與奇鋐份額（每架液冷內容取自橋接表；分歧管、內容調整、滲透率、份額為本模型 Assumed）")
inp("mani", "每架分歧管與快接頭（橋接表未含）", "US$K", 3.0, 6.0, 10.0, "Assumed", "公司：分歧管月產能 7,000 支，約供 3,000–4,000 架（優分析 2026-08-12）→ 每架約 2 支", "分歧管單價：未找到")
for y, lo, b, hi in ((27, 1.0, 1.3, 1.6), (28, 1.0, 1.4, 1.8), (29, 1.0, 1.4, 1.8)):
    inp(f"lcadj_{y}", f"每架液冷內容 ÷ 橋接表 20{y}（Rubin 全液冷、ASIC 液冷；2026＝1）", "x", lo, b, hi, "Assumed",
        ETTF + "：Vera Rubin 無風扇全液冷，液冷延伸到網卡、電源板、光模組，單架散熱價值提升（TrendForce，Interested-party）" if y == 27 else "同上",
        "Morgan Stanley BOM 只計運算匣與交換匣；大和：2027 ASIC 營收超過 GPU（時報 2026-08-13）" if y == 27 else "—")
for y, lo, b, hi in ((26, 0.45, 0.55, 0.65), (27, 0.60, 0.75, 0.85), (28, 0.70, 0.85, 0.95), (29, 0.75, 0.90, 0.95)):
    inp(f"lpen_{y}", f"液冷滲透率 20{y}（部署 GW 中液冷占比）", "%", lo, b, hi, "Assumed",
        ETTF + "：TrendForce AI 晶片液冷滲透率 2025 約 33%、2026 53%、2027 近 60%（以顆數計，Interested-party）" if y == 26 else "同上",
        "以 GW 計高於顆數計（液冷晶片功耗較高）；Trainium 3 先氣冷後液冷（TechNews 2026-04-29）" if y == 26 else "—")
for y, lo, b, hi in ((27, 0.85, 0.95, 1.05), (28, 0.75, 0.90, 1.05), (29, 0.70, 0.85, 1.05)):
    inp(f"rel_{y}", f"奇鋐液冷份額 20{y} ÷ 2026（相對值）", "x", lo, b, hi, "Assumed",
        "凱基 2026-08-05（網址見 Sources）：奇鋐為 CSP 與企業 GPU／ASIC 伺服器氣冷與液冷領導廠（未給份額，Interested-party）" if y == 27 else "同上",
        "對手：雙鴻、酷碼、台達電、Boyd、CoolIT；公司 2027 液冷產能再擴約 50%（優分析 2026-08-12）；份額數字：未找到" if y == 27 else "—")
for y, lo, b, hi in ((27, 0.60, 0.75, 0.90), (28, 0.45, 0.65, 0.85), (29, 0.40, 0.60, 0.80)):
    inp(f"acr_{y}", f"AI 氣冷每 GW 營收 20{y} ÷ 2026（相對值）", "x", lo, b, hi, "Assumed",
        "氣冷轉液冷：液冷滲透率上升，每 GW 氣冷金額下降（本模型推論）" if y == 27 else "同上",
        "Trainium 3、T3 機櫃 2H26 以氣冷為主，T3 液冷 2027 起（凱基 2026-08-05）" if y == 27 else "—")

sec("C. 營收結構錨點（公司揭露比例 × 營收）")
inp("srv_25", "2025 年伺服器與網通 ÷ 營收", "%", None, 0.52, None, "Interested-party", "凱基 2026-08-12（網址見 Sources）：伺服器與網通 2025 年 52%、1H25 48%", "第二來源：未找到（凱基 2026-08-05 另列伺服器 52%＋網通 7%，口徑不一）")
inp("srv_26", "2026 年伺服器與網通 ÷ 營收", "%", 0.66, 0.67, 0.69, "Interested-party", UA2Q + "：1H26 伺服器應用 649.19 億元、占 66.1%（公司）", WR1Q + "：1Q26 伺服器與網通約 66.4%；凱基估 2026 年 69%")
inp("lc_25", "2025 年液冷 ÷ 營收", "%", 0.10, 0.15, 0.20, "Assumed", "未找到（公司稱 2025 年為液冷元年，未揭露比重）", "只影響 FY2025 拆分的顯示，不影響預測")
inp("lc_26", "2026 年液冷 ÷ 營收", "%", 0.28, 0.32, 0.35, "Interested-party", TN0429 + "：澳系券商估今年液冷占營收 30–35%", "凱基：2027 年液冷超過營收 50%（2026 年未給）；公司：液冷占比持續提升（未給數字）")
inp("aiair", "伺服器與網通中非液冷部分屬 AI 的比例（其餘為一般伺服器、網通）", "%", 0.40, 0.60, 0.80, "Assumed", "未找到（公司未拆 AI 與一般伺服器）", "含 AI 伺服器氣冷模組、3D VC、AI 機殼與機櫃、交換器散熱")
for y, lo, b, hi in ((27, 0.0, 0.08, 0.15), (28, 0.0, 0.06, 0.12), (29, 0.0, 0.05, 0.10)):
    inp(f"g_gen_{y}", f"一般伺服器與網通營收年增 20{y}", "%", lo, b, hi, "Assumed", "一般伺服器換機（Intel、AMD 新平台）與交換器" if y == 27 else "同上", "—")
for y, lo, b, hi in ((27, 0.0, 0.06, 0.12), (28, 0.0, 0.05, 0.10), (29, 0.0, 0.04, 0.08)):
    inp(f"g_nai_{y}", f"非伺服器營收年增 20{y}（3C、系統、富世達軸承與鉸鏈）", "%", lo, b, hi, "Assumed",
        "2026 錨點非伺服器約 +11%（Q 頁）；手機與平板 VC 目標營收 5–10%（公司 4Q25 法說，優分析）" if y == 27 else "同上",
        "1H26 機殼 +134%、散熱 +96%（凱基轉述公司）" if y == 27 else "—")

sec("D. 毛利率")
inp("d_lc", "液冷毛利率高於其他營收的幅度", "pt", 0.02, 0.06, 0.10, "Assumed", "公司：氣冷轉液冷顯著提升產品 ASP 與獲利（1Q26 法說，時報 2026-05-14）", "產品別毛利率：未找到；毛利率 FY2025 25.8% → 2Q26 32.6%，同期伺服器比重 52% → 66%")

sec("E. FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設）")
inp("q3_rev", "3Q26 營收（7–9 月月營收合計）", "NT$B", None, "=18.59+19.481+21.018", None, "Verified", TN09 + "：9 月 210.18 億元、前 9 月 1,572.48 億元（1,572.48 − 1H26 981.59 ＝ 590.89）", "7 月 185.9 億（時報 2026-08-13）、8 月 194.81 億（winvest 2026-09）")
inp("q4g", "4Q26 營收季增", "%", 0.08, 0.15, 0.20, "Interested-party", TN1019 + "：高盛估 4Q26 680.57 億元（季增 15%；10–12 月 216.5、225.2、238.9 億）", KGI08 + "：4Q26 626 億元（以 3Q 530.5 億為基礎，季增 18%；3Q 實際 590.9 億）")
inp("q_gm", "3Q、4Q26 毛利率", "%", 0.32, 0.33, 0.345, "Interested-party", "2Q26 32.57%（公司）；工商時報 2026-08-06：法人估 3Q 毛利率朝 33% 靠攏", KGI08 + "：3Q26 33.4%、4Q26 33.5%；高盛 2026 全年 31.7%")
inp("q_opx", "3Q、4Q26 營業費用 ÷ 營收", "%", 0.048, 0.052, 0.056, "Derived", "1Q26 5.26%（毛利 145.97 − 營業利益 120.19）、2Q26 5.13%（公司）", "FY2025 6.1%")
inp("q_nop", "3Q、4Q26 業外收益（每季）", "NT$B", -0.3, 0.2, 0.6, "Derived", "1Q26 −0.04、2Q26 +0.46（稅前 − 營業利益；稅前以 HiStock 稅前淨利率推導）", HIS)
inp("q_nir", "3Q、4Q26 歸屬母公司淨利 ÷ 稅前", "%", 0.66, 0.675, 0.69, "Derived", "1H26 174.83 ÷ 稅前 259.2＝0.675；1Q26 0.661、2Q26 0.687", "含稅率約 23% 與非控制權益（富世達等）")
inp("q_sh", "股數（十億股；EPS 換算）", "十億股", None, "=17.483/44.54", None, "Derived", "1H26 淨利 174.83 億 ÷ EPS 44.54（公司）", "已發行 392,543,007 股（HiStock）")

sec("F. 營業費用、業外、稅與非控制權益、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.5, 0.7, 0.9, "Derived", "FY2026 錨點實測（Q 頁 B26）", "FY2024 營業費用：未找到（FY2025 彈性無法計算）")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 ±2 個百分點")
inp("nop_f", "業外收益（每年）", "NT$B", -1.0, 1.0, 3.0, "Assumed", "FY2025 約 1.2、1H26 約 0.4（推導）", "匯兌損益波動大")
inp("nir_f", "歸屬母公司淨利 ÷ 稅前", "%", 0.66, 0.675, 0.69, "Derived", "同 q_nir", "富世達持股比例變動會影響非控制權益：未建模")
inp("sh_g", "股數年變動", "%", 0.0, 0.0, 0.01, "Assumed", "FY2025 推導 0.390 → 1H26 0.393 十億股（公司 EPS 與淨利）", "可轉債、員工認股：未找到")
inp("mf", "AI 部署係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 AI IT 資本支出約 2 倍", "±20%")

sec("G. 共識、市場資料（只作 Recon 對帳，不作校準）")
inp("b_rev26", "凱基 FY2026 營收（2026-08-12，3Q 前）", "NT$B", None, 213.809, None, "Interested-party", KGI08, "—")
inp("b_q4", "高盛 4Q26 營收（2026-10-19）", "NT$B", None, 68.057, None, "Interested-party", TN1019, "—")
inp("b_eps26gs", "高盛 FY2026 EPS（2026-10-19）", "NT$", None, 101.06, None, "Interested-party", TN1019 + "：2027 154.78、2028 195.18；目標價 4,736（30.6x 2027）", "—")
inp("b_eps26kgi", "凱基 FY2026 EPS（2026-08-12）", "NT$", None, 103.17, None, "Interested-party", KGI08, "美銀 104.95、大和 102.99、花旗 102.7（時報 2026-08-13）")
inp("b_rev27", "凱基 FY2027 營收（2026-08-12）", "NT$B", None, 311.893, None, "Interested-party", KGI08, "花旗：2027 營收年增 50%（經濟日報 2026-09-19）")
inp("b_eps27gs", "高盛 FY2027 EPS（2026-10-19）", "NT$", None, 154.78, None, "Interested-party", TN1019, "—")
inp("b_eps27kgi", "凱基 FY2027 EPS（2026-08-12）", "NT$", None, 160.03, None, "Interested-party", KGI08, "—")
inp("b_eps27bofa", "美銀 FY2027 EPS（2026-08-13 前後）", "NT$", None, 143.99, None, "Interested-party", WR2Q + "：美銀 2026 104.95、2027 143.99、2028 178.74", "—")
inp("b_eps27dw", "大和 FY2027 EPS（2026-08-13 前後）", "NT$", None, 139.53, None, "Interested-party", WR2Q + "：大和 2026 102.99、2027 139.53、2028 183.06", "—")
inp("b_eps27citi", "花旗 FY2027 EPS（2026-08-13 前後）", "NT$", None, 153.21, None, "Interested-party", "時報 2026-08-13（網址見 Sources）：花旗 2027 153.21、2028 185.53", "花旗 2026-09-19 再上修 2027 獲利 13%（經濟日報；絕對值未找到）")
inp("b_eps28gs", "高盛 FY2028 EPS（2026-10-19）", "NT$", None, 195.18, None, "Interested-party", TN1019, "—")
inp("b_gm27", "花旗 FY2027 毛利率（2026-09-19）", "%", None, 0.349, None, "Interested-party", "經濟日報 2026-09-19（花旗）", "高盛 33.4%、凱基 33.8%")
inp("b_lc27", "凱基：2027 年液冷 ÷ 營收（超過）", "%", None, 0.50, None, "Interested-party", KGI08, "大和：2027 液冷滲透率超過 50%")
inp("tf_pen26", "TrendForce：AI 晶片液冷滲透率 2026（顆數）", "%", None, 0.53, None, "Interested-party", ETTF, "—")
inp("tf_pen27", "TrendForce：AI 晶片液冷滲透率 2027（顆數）", "%", None, 0.60, None, "Interested-party", ETTF, "「近 60%」")

sec("H. 價格與估值")
inp("px", "股價（2026-10-22 收盤）", "NT$", None, 3145, None, "Interested-party", "HiStock、鉅亨（10-22 收盤 3,145，−2.78%）", "FinLab（10-24 更新）仍為 10-22 收盤；10-23 收盤：未找到")
inp("pe_ttm", "近 12 月 P/E（10-22）", "x", None, 41.92, None, "Interested-party", "HiStock（10-22）", "財報狗 41.98")

sec("I. 目標價情境權重（Val 第 0 節用；Andy 可改）")
for s_, w_, nm_ in (("A", 0.3, "順利"), ("B", 0.5, "平平"), ("C", 0.2, "不如預期")):
    inp(f"tp_w{s_}", f"目標價情境權重：情境 {s_}（{nm_}）", "權重", None, w_, None, "Assumed",
        "Andy 可改（2026-10-25 目標價規格預設 A／B／C＝0.3／0.5／0.2）", "三個權重合計應為 1；v0.2 起用於 DCF 與 Comparable 的情境加權（Val 第 0 節）")

# ---------------- v0.2 設定與新增輸入（附加在 In 既有輸入之後，不位移既有列） ----------------
V2 = {'name': '奇鋐',
 'unit': 'NT$B',
 'shunit': '十億股',
 'cur': 'NT$',
 'shg': 'sh_g',
 'epsfmt': '0.00',
 'pxfmt': '#,##0',
 'keys': {'rev': 'rev', 'ni': 'ni', 'sh': 'sh_', 'eps': 'eps'},
 'roe': ('=AVERAGE(0.2688,0.3099,0.4986)', 'stockanalysis 比率頁（Return on Equity）：FY2023 26.88%、FY2024 30.99%、FY2025 49.86%', '單一來源；近 12 月 66.78%（同站，2026-10-12）'),
 'cons_tp': (4130, 'stockanalysis forecast 頁（S&P Global）17 位分析師平均 4,130（中位數 4,250），頁面時間 2026-10-01', '個股頁（2026-10-05）4,200.35；第二來源：未找到'),
 'cons_feps': ('=3585/27.04', 'stockanalysis 個股頁 2026-10-05：價格 3,585 ÷ Forward PE 27.04'),
 'peers': [('雙鴻（Auras）', '3324.TWO', 17.27, '2026-08-28'), ('健策（Jentech）', '3653.TW', 59.27, '2026-08-27'), ('Vertiv', 'VRT', 24.53, '2026-10-23'), ('台達電（Delta）', '2308.TW', 30.51, '2026-10-22')],
 'peer_note': '散熱與液冷同業：雙鴻、健策、Vertiv、台達電（券商常用比較組）。英維克（002837.SZ）為 A 股，依規則不列。非 A 股公司的同業名單不放中國 A 股（A 股有結構性溢價；Andy 2026-10-25 11:53）',
 'fwd_w': (0.16666666666666666, 0.8333333333333334),
 'fwd_note': '約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）',
 'pe10': None,
 'px_note': '現價＝In 頁 px（2026-10-22 收盤 3,145；10-23 收盤未找到）',
 'method': 'DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年平均約 36%（再投資率約 11%），FY2027–29 驅動因子、FY2030–32 收斂法＋永續；ROE 高使再投資率低，DCF 偏高估。Comparable：雙鴻、健策、Vertiv、台達電前瞻 P/E 中位數；同業倍數分散（約 17–59x），雙鴻、健策價格為 8 '
           '月底，偏誤方向不明確。情境對應：A＝順利（顆數基準＋液冷滲透率高值）、B＝平平（顆數 B＋基準）、C＝不如預期（顆數 C＋滲透率低值）。',
 'readme_val': '估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），FY2027–FY2032 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，推到 12 個月後）＋0.5 × Comparable 目標價（雙鴻、健策、Vertiv、台達電 前瞻 P/E 中位數 × 12 個月後前瞻 '
               'EPS）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解（EPS 因素、倍數因素）、DCF 價格帶、十年 P/E 分位檢查（只作檢查）。',
 'readme_add': ['v0.2（2026-10-25）：(1) PL 頁最下方新增延伸表，三情境延伸到 FY2032；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt；淨利率由 FY2029 收斂到 A＝FY2029、B＝FY2029 '
                '與十年平均的中點、C＝十年平均；十年平均淨利率未找到，依工作單以 FY2029 × 0.8 替代，Assumed；PL 延伸表灰底）。(2) Val 第 0 節目標價＝DCF 與 Comparable 加權（tp_w_dcf＝0.5）；v0.1.1 的「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留，不再是目標價。(3) Sens '
                '新增目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.4／0.8）。(4) Recon 新增目標價對帳。(5) 現價：10-23 收盤未找到（Yahoo 股市、HiStock、鉅亨於 2026-10-25 仍為 10-22 或更早），依工作單沿用 2026-10-22 收盤 3,145。上游台達電模型改為 v0.2，本模型讀取的 '
                'Units 液冷線數值不變。 新增內容一律放在各頁既有內容之後，FY2025–FY2029 既有數值不變。',
                '情境對應：A＝順利（顆數基準、FY2027 起液冷滲透率高值）、B＝平平（CY2028 顆數 ×0.8、CY2029 ×0.85，滲透率基準）、C＝不如預期（CY2028–29 ×0.6、滲透率低值）。v0.2.1 起液冷滲透率依情境字母（Andy 2026-10-25 13:13，與雙鴻模型同值），FY2027 三情境只差滲透率，顆數差異從 CY2028 起。',
                '先驗：本版未做先驗，無 Andy 合理 P/E 對照。']}
sec("J. 目標價（DCF 與 Comparable 加權）與收斂法參數（v0.2 新增；黃底＝Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（DCF 永續；情境 B 的 FY2032 營收年增收斂目標）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "Sens 頁 ±1pt（只動 DCF，PL 收斂目標不隨之變動）")
inp("tp_roe", "ROE（近三年平均）", "%", None, V2["roe"][0], None, "Derived", V2["roe"][1], V2["roe"][2])
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 改為 0.5：DCF／Comparable 各半）", "Sens 頁 0.4／0.8")
inp("fade_gA", "情境 A（順利）：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂，FY2032 達到")
inp("fade_gB", "情境 B（平平）：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C（不如預期）：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("nim10_f", "十年平均淨利率的替代：最後一年（FY2029）淨利率 ×（本值）", "x", None, 0.8, None, "Assumed", "十年（2016–2025）平均淨利率：未找到（免費資料商只到近五年）；依工作單以最後一年 × 0.8 替代", "淨利率收斂目標：A＝FY2029、B＝FY2029 與替代值的中點、C＝替代值（PL 延伸表）")
inp("cons_tp", "市場共識目標價（賣方平均）", V2["cur"], None, V2["cons_tp"][0], None, "Interested-party", V2["cons_tp"][1], V2["cons_tp"][2])
inp("cons_feps", "共識前瞻 12 月 EPS（＝價格 ÷ 前瞻 P/E）", V2["cur"], None, V2["cons_feps"][0], None, "Interested-party", V2["cons_feps"][1], "賣方共識，只作對帳與差距拆解，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣；股數為十億股）。藍字＝輸入；黑字＝公式。實績取自奇鋐法說（經媒體轉述）、凱基報告與月營收"
Q["A2"] = "FY2026 錨點＝1H26 實績＋3Q26 營收實績（月營收）＋3Q、4Q 假設（毛利率、營業費用率、業外、淨利 ÷ 稅前；4Q 營收＝3Q × (1＋季增)）。3Q26 財報（預計 11 月中）尚未公布。稅前利益以 HiStock 各季稅前淨利率 × 營收推導（公司稅前金額原文：未找到）；FY2025 各季營收以年增率回推"
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（營收實績）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：" + UA4Q + "；1H26：" + UA2Q
S2 = "凱基 2026-08-12（" + KGI08 + "）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
# FY2025 各季營收：4Q25 477.75（公司）；3Q25＝477.75 ÷ 1.227（季增 22.7%）；1H25＝1H26 981.59 ÷ 1.8546；1Q25＝1Q26 490.38 ÷ 2.1017
PT25 = "=(49.038/2.1017)*0.198+(98.159/1.8546-49.038/2.1017)*0.1992+(47.775/1.227)*0.211+47.775*0.2104"
qrow(4, "營收", {"B": 139.639, "C": 98.159, "D": "=q3_rev", "E": "=D4*(1+q4g)"}, tag="Verified／Interested-party（4Q 季增）")
qrow(5, "毛利", {"B": 36.034, "C": "=14.597+49.121*0.3257", "D": "=D4*q_gm", "E": "=E4*q_gm"}, tag="Verified／Derived（2Q26＝營收 × 毛利率）", s2="FY2025 毛利率 25.81%（公司）；1Q26 毛利 145.97 億（時報 2026-05-14）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2025 25.8%；1H26 31.17%（公司）")
qrow(7, "營業費用", {"B": "=B5-B8", "C": "=C5-C8", "D": "=D4*q_opx", "E": "=E4*q_opx"}, tag="Derived（毛利 − 營業利益）／Assumed")
qrow(8, "營業利益", {"B": "=139.639*0.1973", "C": "=12.019+13.478", "D": "=D5-D7", "E": "=E5-E7"}, tag="Derived（FY2025 營益率 19.73%）／Verified", s2="1Q26 120.19 億（時報）；2Q26 134.78 億（凱基）")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2025 19.7%；1H26 25.97%（公司）")
qrow(10, "業外收益", {"B": "=B11-B8", "C": "=C11-C8", "D": "=q_nop", "E": "=q_nop"}, tag="Derived")
qrow(11, "稅前利益", {"B": PT25, "C": "=49.038*0.2443+49.121*0.2837", "D": "=D8+D10", "E": "=E8+E10"}, tag="Derived（稅前淨利率 × 營收）", s1=HIS, s2="各季稅前淨利率：1Q25 19.80%、2Q25 19.92%、3Q25 21.10%、4Q25 21.04%、1Q26 24.43%、2Q26 28.37%")
qrow(12, "歸屬母公司淨利", {"B": 19.186, "C": 17.483, "D": "=D11*q_nir", "E": "=E11*q_nir"}, tag="Verified／Derived（3Q、4Q）")
qrow(13, "淨利 ÷ 稅前（含稅與非控制權益）", {c: f"={c}12/{c}11" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(14, "股數（十億股）", {"B": "=19.186/49.17", "C": "=q_sh", "D": "=q_sh", "E": "=q_sh"}, ann="=q_sh", tag="Derived", s2="淨利 ÷ EPS")
qrow(15, "EPS", {c: f"={c}12/{c}14" for c in "BCDE"}, fmt="0.00", ann="=F12/F14", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(16, "  公司揭露 EPS", {"B": 49.17, "C": 44.54}, fmt="0.00", ann="", tag="Verified", s2="1Q26 20.17、2Q26 24.37；4Q25 16.92")
SPL = [  # (列, 標籤, FY2025 公式, FY2026 公式)
    (17, "液冷營收（＝營收 × 液冷占比）", "=B4*lc_25", "=F4*lc_26"),
    (18, "AI 氣冷與機構（＝營收 ×（伺服器與網通 − 液冷）× AI 比例）", "=B4*(srv_25-lc_25)*aiair", "=F4*(srv_26-lc_26)*aiair"),
    (19, "一般伺服器與網通（＝營收 ×（伺服器與網通 − 液冷）×（1 − AI 比例））", "=B4*(srv_25-lc_25)*(1-aiair)", "=F4*(srv_26-lc_26)*(1-aiair)"),
    (20, "非伺服器（＝營收 ×（1 − 伺服器與網通）；3C、系統、富世達）", "=B4*(1-srv_25)", "=F4*(1-srv_26)"),
    (21, "AI 氣冷（Sens：AI 比例 低）", None, "=F4*(srv_26-lc_26)*aiair_lo"),
    (22, "一般伺服器（Sens：AI 比例 低）", None, "=F4*(srv_26-lc_26)*(1-aiair_lo)"),
    (23, "AI 氣冷（Sens：AI 比例 高）", None, "=F4*(srv_26-lc_26)*aiair_hi"),
    (24, "一般伺服器（Sens：AI 比例 高）", None, "=F4*(srv_26-lc_26)*(1-aiair_hi)"),
    (25, "液冷（Sens：2026 液冷占比 低）", None, "=F4*lc_26_lo"),
    (26, "AI 氣冷（Sens：液冷占比 低）", None, "=F4*(srv_26-lc_26_lo)*aiair"),
    (27, "一般伺服器（Sens：液冷占比 低）", None, "=F4*(srv_26-lc_26_lo)*(1-aiair)"),
    (28, "液冷（Sens：2026 液冷占比 高）", None, "=F4*lc_26_hi"),
    (29, "AI 氣冷（Sens：液冷占比 高）", None, "=F4*(srv_26-lc_26_hi)*aiair"),
    (30, "一般伺服器（Sens：液冷占比 高）", None, "=F4*(srv_26-lc_26_hi)*(1-aiair)")]
for rr, lab, fb, ff in SPL:
    qrow(rr, lab, {"B": fb} if fb else {}, ann=ff, tag="Derived／Interested-party" if rr < 21 else "Derived", s1="公司產品比重（In 頁 C 節）" if rr < 21 else "Sens 用", s2="液冷、AI 與一般伺服器的季度金額未揭露：未找到" if rr == 17 else "—")
Q.cell(32, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(33, 1, "FY2026 營收年增"); Q["B33"] = "=F4/B4-1"
Q.cell(34, 1, "FY2026 營業費用年增"); Q["B34"] = "=F7/B7-1"
Q.cell(35, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B35"] = "=B34/B33"
Q.cell(36, 1, "FY2026 營收 ÷ 凱基（08-12，3Q 前）− 1"); Q["B36"] = "=F4/b_rev26-1"
Q.cell(37, 1, "FY2026 EPS ÷ 高盛（10-19）− 1"); Q["B37"] = "=F15/b_eps26gs-1"
Q.cell(38, 1, "非伺服器營收 FY2026 年增"); Q["B38"] = "=F20/B20-1"
for rr in range(33, 39): Q[f"B{rr}"].number_format = "0.0%" if rr != 35 else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("nop", 10), ("pt", 11), ("ni", 12), ("sh", 14), ("eps", 15),
               ("lc", 17), ("air", 18), ("gen", 19), ("nai", 20), ("airL", 21), ("genL", 22), ("airH", 23), ("genH", 24),
               ("lcLL", 25), ("airLL", 26), ("genLL", 27), ("lcLH", 28), ("airLH", 29), ("genLH", 30)):
    if rr <= 20: defname(f"a25_{nm}", f"Q!$B${rr}")
    defname(f"a26_{nm}", f"Q!$F${rr}")
Q.column_dimensions["A"].width = 64
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 90; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# 液冷滲透率屬行業情境（技術替代），依情境字母（Andy 2026-10-25 13:13；v0.2.1）：FY2027 起 A＝高值、B＝基準、C＝低值，FY2026 三情境皆基準。
# 與雙鴻模型（auras/）同一做法；數值單一來源＝本頁 In 的 lpen_26～lpen_29（低／基準／高），雙鴻讀此三欄。
for y in (26, 27, 28, 29):
    for s_, suf in (("A", "_hi"), ("B", ""), ("C", "_lo")):
        wb.defined_names[f"lpen{s_}_{y}"] = DefinedName(f"lpen{s_}_{y}", attr_text=wb.defined_names[f"lpen_{y}" + ("" if y == 26 else suf)].attr_text)

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", rel="rel", lcadj="lcadj", lpen=None, acr="acr", mf="1", gg="g_gen", gn="g_nai", d="d_lc", opxa="0", mani="mani",
            lc="lc", air="air", gen="gen")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    P["lpen"] = P["lpen"] or f"lpen{s}"   # 液冷滲透率依情境字母（v0.2.1）
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    m = P["mani"]
    sec_("A. AI 部署與液冷市場（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029）")
    row("itc", "AI IT 資本支出 × 情境顆數係數", "US$B", {c_(y): f"=itc_{y % 100}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("itpg", "每 GW IT 資本支出", "US$B", {"D": "=itpg_gb*(1-vr26)+itpg_vr*vr26", **{c_(y): "=itpg_vr" for y in BU[1:]}})
    row("gw", "部署 GW（IT 電力）", "GW", {c_(y): f"={c_(y)}{R['itc']}/{c_(y)}{R['itpg']}" for y in BU}, fmt="0.0")
    row("lpen", "液冷滲透率（部署 GW 中液冷占比）", "%", {c_(y): f"={P['lpen']}_{y % 100}" for y in BU}, fmt="0.0%")
    row("cpg", "每 GW 液冷內容（機架內液冷＋分歧管）", "US$M/GW",
        {"D": f"=((1-vr26)*racks_gb*(cool_gb+{m})+vr26*racks_vr*(cool_vr+{m}))/1000",
         **{c_(y): f"=racks_vr*(cool_vr+{m})*{P['lcadj']}_{y % 100}/1000" for y in BU[1:]}}, fmt="#,##0",
        note="橋接表每架機架內液冷 × 每 GW 機架數；FY2027 起 × 內容調整（Rubin 全液冷、ASIC 液冷）")
    row("lmk", "AI 液冷市場（奇鋐可做部分；美元）", "US$B", {c_(y): f"={c_(y)}{R['gw']}*{c_(y)}{R['lpen']}*{c_(y)}{R['cpg']}/1000" for y in BU})
    row("lmn", "AI 液冷市場（新台幣）", "NT$B", {c_(y): f"={c_(y)}{R['lmk']}*fx" for y in BU}, fmt="#,##0.0")
    row("sh0", "奇鋐有效份額（FY2026 校準）＝ 液冷營收 ÷ 市場", "%", {"D": f"=a26_{P['lc']}/D{R['lmn']}"}, fmt="0.0%",
        note="Derived：分子為公司液冷占比（券商說法）× 營收；分母為本模型的液冷市場（不含 CDU、機架外冷卻）")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "奇鋐有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("lc", "液冷營收", "NT$B", {"C": "=a25_lc", **{c_(y): f"={c_(y)}{R['lmn']}*{c_(y)}{R['sh']}" for y in BU}}, fmt="#,##0.0")
    sec_("B. AI 氣冷、一般伺服器、非伺服器與合計")
    row("apg0", "AI 氣冷與機構每 GW 營收（FY2026 校準）", "NT$B/GW", {"D": f"=a26_{P['air']}/D{R['gw']}"}, fmt="0.000")
    row("arel", "AI 氣冷每 GW 相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['acr']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("air", "AI 氣冷與機構營收", "NT$B", {"C": "=a25_air", **{c_(y): f"={c_(y)}{R['gw']}*$D${R['apg0']}*{c_(y)}{R['arel']}" for y in BU}}, fmt="#,##0.0")
    row("gen", "一般伺服器與網通營收", "NT$B", {"C": "=a25_gen", "D": f"=a26_{P['gen']}", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gg']}_{y % 100})" for y in BU[1:]}}, fmt="#,##0.0")
    row("nai", "非伺服器營收（3C、系統、富世達）", "NT$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']}_{y % 100})" for y in BU[1:]}}, fmt="#,##0.0")
    row("rev", "營收", "NT$B", {c_(y): f"={c_(y)}{R['lc']}+{c_(y)}{R['air']}+{c_(y)}{R['gen']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.0")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("lcsh", "液冷 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['lc']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("aish", "AI（液冷＋氣冷與機構）÷ 營收", "%", {c_(y): f"=({c_(y)}{R['lc']}+{c_(y)}{R['air']})/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("pgw", "每 GW 奇鋐 AI 營收（液冷＋氣冷與機構）", "US$M/GW", {c_(y): f"=({c_(y)}{R['lc']}+{c_(y)}{R['air']})/fx/{c_(y)}{R['gw']}*1000" for y in BU}, fmt="#,##0")
    sec_("C. 損益")
    row("gmn", "其他營收毛利率 ＝ 錨點毛利率 − 差距 × 液冷占比（FY2026 校準）", "%", {"D": f"=a26_gp/D{R['rev']}-{P['d']}*D{R['lcsh']}"}, fmt="0.0%",
        note="Derived：讓 FY2026 毛利率等於錨點；液冷毛利率＝其他 + 差距。FY2027 起毛利率隨液冷占比變動")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"=$D${R['gmn']}*{c_(y)}{R['rev']}+{P['d']}*{c_(y)}{R['lc']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in BU[1:]}}, fmt="#,##0.0")
    row("opx", "營業費用", "NT$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("oi", "營業利益", "NT$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.0")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("nop", "業外收益", "NT$B", {"C": "=a25_nop", "D": "=a26_nop", **{c_(y): "=nop_f" for y in BU[1:]}}, fmt="#,##0.0")
    row("nir", "歸屬母公司淨利 ÷ 稅前", "%", {"C": "=a25_ni/(a25_oi+a25_nop)", "D": "=a26_ni/(a26_oi+a26_nop)", **{c_(y): "=nir_f" for y in BU[1:]}}, fmt="0.0%")
    row("ni", "歸屬母公司淨利", "NT$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}+{c_(y)}{R['nop']})*{c_(y)}{R['nir']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.000")
    row("eps", "EPS", "NT$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 奇鋐損益（十億新台幣；EPS 新台幣；歸屬母公司）。FY2025A 實績、FY2026 錨點（1H26 實績＋3Q26 營收實績＋假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026 三情境相同；FY2027 起液冷滲透率依情境字母（A 高、B 基準、C 低；v0.2.1），顆數係數從 CY2028 開始"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準顆數（係數 1.0）＋液冷滲透率高值（FY2027 起）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85＋液冷滲透率基準", "C": "情境 C：CY2028–29 顆數 ×0.6＋液冷滲透率低值（FY2027 起）"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("lc", "  液冷營收", "NT$B", "#,##0.0"), ("air", "  AI 氣冷與機構營收", "NT$B", "#,##0.0"),
                                ("gen", "  一般伺服器與網通營收", "NT$B", "#,##0.0"), ("nai", "  非伺服器營收", "NT$B", "#,##0.0"), ("lcsh", "  液冷 ÷ 營收", "%", "0.0%"),
                                ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"),
                                ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("eps", "EPS", "NT$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收 − 錨點；FY2026 模型液冷 − 錨點；FY2026 毛利 − 錨點").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：液冷差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['lc']}-a26_lc"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（NT$B）"); PL[f"D{sr}"] = f"=D{PB['A']['gmn']}*D{PB['A']['rev']}+d_lc*D{PB['A']['lc']}-a26_gp"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：液冷內容、滲透率、份額 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — AI 液冷市場（奇鋐可做部分：冷板、分歧管、快接頭）、奇鋐份額與每 GW 營收。部署 GW＝AI IT 資本支出（markets/network）÷ 每 GW IT 資本支出（Tokenomics，經橋接表）；每架液冷取自橋接表"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("gw", "部署 GW（情境 A）", "GW", lambda y: f"=PL!{FC[y]}{A_['gw']}", "0.0", "AI IT 資本支出 ÷ 每 GW IT 資本支出"),
        ("lpen", "液冷滲透率（部署 GW）", "%", lambda y: f"=PL!{FC[y]}{A_['lpen']}", "0.0%", "In 頁 B 節（Assumed）；情境 A 取高值（FY2027 起）"),
        ("cpg", "每 GW 液冷內容", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['cpg']}", "#,##0", "橋接表機架內液冷＋分歧管 × 每 GW 機架數 × 內容調整"),
        ("crk", "每架液冷內容（以 VR200 每 GW 架數換算）", "US$K/架", lambda y: f"=PL!{FC[y]}{A_['cpg']}*1000/racks_vr", "#,##0", "對照：橋接表 VR200 機架內液冷 US$56K＋分歧管"),
        ("lmk", "AI 液冷市場（奇鋐可做部分）", "US$B", lambda y: f"=PL!{FC[y]}{A_['lmk']}", "#,##0.0", "不含 CDU、機架外冷卻與液對氣"),
        ("sh", "奇鋐有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以券商液冷占比校準"),
        ("lc", "奇鋐液冷營收", "US$B", lambda y: f"=PL!{FC[y]}{A_['lc']}/fx", "#,##0.00", ""),
        ("air", "奇鋐 AI 氣冷與機構營收", "US$B", lambda y: f"=PL!{FC[y]}{A_['air']}/fx", "#,##0.00", "FY2026 以伺服器比重與 AI 比例校準"),
        ("pgw", "每 GW 奇鋐 AI 營收（液冷＋氣冷與機構）", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgw']}", "#,##0", ""),
        ("oth", "其他廠商液冷（雙鴻、酷碼、台達電、Boyd、CoolIT 等）＝ 市場 − 奇鋐", "US$B", lambda y: f"=PL!{FC[y]}{A_['lmk']}-PL!{FC[y]}{A_['lc']}/fx", "#,##0.0", "")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：VR200 每 GW 機架數較少（6,764 → 4,052 架），橋接表的每 GW 液冷內容下降；FY2027 起的成長來自液冷滲透率與每架內容調整（Rubin 全液冷、ASIC 液冷），兩者與份額相對值都是 Assumed，是最需要每季檢查的輸入。").font = BOLD
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」與十年 P/E 規則值（年底落後 P/E，基準為 FY2027 EPS）"
V["A2"] = "循環位置：奇鋐不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照台達電、NVIDIA、Arista；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-22 收盤）"); V["B6"] = "=px"; V["B6"].number_format = "#,##0"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "未來 12 月 EPS（模型 A）≈ ¼ FY2026＋¾ FY2027"); V["B11"] = f"=0.25*PL!D{SUMR[('A', 'eps')]}+0.75*PL!E{SUMR[('A', 'eps')]}"; V["B11"].number_format = "0.00"
V.cell(12, 1, "  現價 ÷ 未來 12 月 EPS（模型 A）"); V["B12"] = "=px/B11"; V["B12"].number_format = "0.0"
V.cell(13, 1, "近 12 月 P/E（HiStock，10-22）"); V["B13"] = "=pe_ttm"; V["B13"].number_format = "0.0"
V["E8"] = "FY2027 三情境只差液冷滲透率（A 高、B 基準、C 低；v0.2.1）；顆數係數從 CY2028 開始"
V.cell(15, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, 16, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 17):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(21, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 22, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 23 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{17 + j}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(26, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, 27, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數", "對照：十年 75 分位"])
vr = 28
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, "=pe_mid").number_format = "0.0"; V.cell(vr, 5, "=pe_trough").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。本版未做先驗，無 Andy 合理 P/E 對照。")
vr += 1
assert vr < 40, vr
# 十年年底 P/E：Goodinfo 年底收盤 ÷ 當年 EPS（2016–2025）；stockanalysis 年底 P/E（2021–2025，只作核對）
GI_PX = [25.1, 32.4, 25.25, 41.65, 66.1, 88.2, 112, 336.5, 623, 1510]
GI_EPS = [2.43, 2.51, 2.01, 2.71, 5.42, 8.21, 11.78, 14.11, 21.21, 49.17]
PE_GI = [f"={p}/{e}" for p, e in zip(GI_PX, GI_EPS)]
PE_SA = [None, None, None, None, None, 10.74, 9.51, 24.32, 29.54, 30.55]
block_end = VR.add_rule_block(V, 40, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "Goodinfo（年底收盤 ÷ 當年 EPS）", "vals": PE_GI, "tag": "Derived（資料商年底收盤與 EPS；2016–2025 共 10 年）"},
    {"label": "stockanalysis（年底 P/E）", "vals": PE_SA, "tag": "Interested-party（資料商；只有 2021–2025，只作核對、不採用）", "use": False,
     "note": "stockanalysis 2021–2025 與 Goodinfo 推導差距在 0.5x 內；2016–2020 年：未找到第二來源"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位。採用 Goodinfo 年底收盤 ÷ 當年 EPS（10 年）；stockanalysis 只有 2021–2025 且全為 AI 前後高倍數期，只作核對。2016–2022 年奇鋐為筆電、伺服器氣冷散熱廠（約 10–15x），AI 前後的倍數口徑不同，規則值可能低估 AI 期的合理倍數。")
rule_rows = [rw for rw in range(40, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# ---------------- Val 第 0 節：目標價與上檔空間（v0.1.1 新增；放在既有空白列，不位移既有列） ----------------
TP0 = V.max_row + 2; TPR = {}
assert all(V.cell(rw, c).value is None for rw in range(TP0 - 1, TP0 + 9) for c in range(1, 9)), TP0
assert V["A3"].value is None
V["A3"] = f"目標價與上檔空間：見第 0 節（本頁第 {TP0} 列；v0.1.1 新增，為了不位移既有列而放在空白列）"
V.cell(TP0, 1, "舊法：規則值情境加權（參考；v0.1.1 的「0. 目標價與上檔空間」改名，不再是目標價；權重為 In 頁 tp_wA／tp_wB／tp_wC）").font = BOLD
for k_, (key_, lab_) in enumerate((("tp", "舊法隱含價（情境加權，參考）"), ("up", "舊法隱含價相對現價（÷ 現價 − 1）"),
        ("A", "情境 A 隱含價（樂觀）"), ("B", "情境 B 隱含價（基準）"), ("C", "情境 C 隱含價（下檔）"),
        ("w", "情境權重 A／B／C"), ("m", "舊法說明")), TP0 + 1):
    V.cell(k_, 1, lab_); TPR[key_] = k_
V.cell(TPR["tp"], 2, f"=tp_wA*B{TPR['A']}+tp_wB*B{TPR['B']}+tp_wC*B{TPR['C']}").number_format = "#,##0"
V.cell(TPR["tp"], 5, "＝權重 A × 情境 A 隱含價＋權重 B × 情境 B 隱含價＋權重 C × 情境 C 隱含價（NT$）")
V.cell(TPR["up"], 2, f"=B{TPR['tp']}/px-1").number_format = "0.0%"
V.cell(TPR["up"], 5, "相對「" + V["A6"].value + "」")
for s_ in SC:
    V.cell(TPR[s_], 2, f"=E{VAL[s_]}").number_format = "#,##0"
    V.cell(TPR[s_], 5, f"＝規則值隱含股價 E{VAL[s_]}（{SCN_T[s_]}）")
for j_, s_ in enumerate(SC):
    V.cell(TPR["w"], 2 + j_, f"=tp_w{s_}").number_format = "0%"
V.cell(TPR["w"], 5, "In 頁輸入（Assumed，Andy 可改）：B 欄＝情境 A、C 欄＝情境 B、D 欄＝情境 C")
V.cell(TPR["m"], 2, "P/E 規則值：各情境 FY2027 EPS × 2016–2025 年底 P/E 依循環位置取分位（見頂＝25 分位、中段＝中位數、谷底＝75 分位；Goodinfo 年底收盤 ÷ 當年 EPS，stockanalysis 核對）。隱含價約為 FY2027 年底價、未折現，作 12 個月目標價的近似；情境 A／B／C 即本模型三個顆數情境（A 最高、C 最低）。偏誤：十年倍數含 2016–2022 年氣冷散熱期低倍數（約 10–15x），目標價可能偏低。")
for c_ in range(1, 6):
    V.cell(TPR["tp"], c_).fill = KEY; V.cell(TPR["up"], c_).fill = KEY
V.cell(TPR["tp"], 1).font = BOLD
defname("old_rule_px", f"Val!$B${TPR['tp']}"); defname("old_rule_up", f"Val!$B${TPR['up']}")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A，含 FY2027 起液冷滲透率高值；有效份額、AI 氣冷每 GW、其他營收毛利率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、液冷營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 液冷營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 液冷營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("份額流失（低）", "份額相對值 2027 0.85、2028 0.75、2029 0.70（雙鴻、酷碼、台達電追趕）", dict(rel="RLO")),
         ("份額上升（高）", "份額相對值 1.05", dict(rel="RHI")),
         ("每架液冷內容 低", "內容調整 1.0（只有橋接表的機架內液冷）", dict(lcadj="CALO")), ("每架液冷內容 高", "內容調整 2027 1.6、2028–29 1.8", dict(lcadj="CAHI")),
         ("液冷滲透率 低", "2026 45%、2027 60%、2028 70%、2029 75%（相對情境 A 的高值）", dict(lpen="LPLO")), ("液冷滲透率 高", "2026 65%、2027 85%、2028–29 95%（FY2027 起與情境 A 相同，只差 FY2026 份額校準）", dict(lpen="LPHI")),
         ("AI 氣冷每 GW 低", "相對值 0.60／0.45／0.40", dict(acr="ACLO")), ("AI 氣冷每 GW 高", "相對值 0.90／0.85／0.80", dict(acr="ACHI")),
         ("AI 部署 ×0.8", "FY2027 起 AI IT 資本支出 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("非 AI 成長 低", "一般伺服器與非伺服器 2027–29 0%", dict(gg="GGLO", gn="GNLO")), ("非 AI 成長 高", "一般伺服器 +15／12／10%、非伺服器 +12／10／8%", dict(gg="GGHI", gn="GNHI")),
         ("液冷毛利率差距 2pt", "液冷毛利率只高 2 個百分點", dict(d="d_lc_lo")), ("液冷毛利率差距 10pt", "高 10 個百分點", dict(d="d_lc_hi")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi")),
         ("2026 液冷占比 28%（讀法）", "FY2026 液冷＝營收 × 28%；份額分子變小、氣冷變大", dict(lc="lcLL", air="airLL", gen="genLL")),
         ("2026 液冷占比 35%（讀法）", "FY2026 液冷＝營收 × 35%", dict(lc="lcLH", air="airLH", gen="genLH")),
         ("AI 比例 40%（讀法）", "伺服器非液冷中屬 AI 的比例 40%", dict(air="airL", gen="genL")),
         ("AI 比例 80%（讀法）", "80%", dict(air="airH", gen="genH"))]
for nm, base, suf, yrs in (("RLO", "rel", "_lo", (27, 28, 29)), ("RHI", "rel", "_hi", (27, 28, 29)), ("CALO", "lcadj", "_lo", (27, 28, 29)), ("CAHI", "lcadj", "_hi", (27, 28, 29)),
                           ("LPLO", "lpen", "_lo", (26, 27, 28, 29)), ("LPHI", "lpen", "_hi", (26, 27, 28, 29)), ("ACLO", "acr", "_lo", (27, 28, 29)), ("ACHI", "acr", "_hi", (27, 28, 29)),
                           ("GGLO", "g_gen", "_lo", (27, 28, 29)), ("GGHI", "g_gen", "_hi", (27, 28, 29)), ("GNLO", "g_nai", "_lo", (27, 28, 29)), ("GNHI", "g_nai", "_hi", (27, 28, 29))):
    for y in yrs:
        target = f"{nm}_{y}"
        wb.defined_names[target] = DefinedName(target, attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "lc", 2027, "#,##0.0"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "lc", 2028, "#,##0.0"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 B（全部）", "CY2028 ×0.8、CY2029 ×0.85＋液冷滲透率基準", PB["B"], "PL!")
put_case(nn + 1, "情境 C（全部）", "CY2028–29 ×0.6＋液冷滲透率低值", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases); SENS_NOTE = nn + 3
SE.cell(nn + 3, 1, "讀法：FY2027 三情境只差液冷滲透率（A 高、B 基準、C 低）；顆數情境 B、C 的差異從 FY2028 起。每個案例的 FY2026 有效份額與 AI 氣冷每 GW 依該案例重算，所以 FY2026 恆等於錨點。「液冷滲透率」案例的 FY2026 份額會反向調整（滲透率低 → 份額高），所以只影響 FY2027 起的成長路徑。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 券商估計、研調滲透率與台達電液冷。只作事後對照，不為對上數字而調參數。券商數字多經媒體轉述；付費共識未找到"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 券商估計（校準未使用）").font = BOLD; rr_ += 1
RECON_A = rr_
rec("FY2026 營收 vs 凱基（08-12，3Q 前）", "=b_rev26", pl("rev", 2026), "凱基 2026-08-12", "凱基 3Q 估 530.5 億，實際 590.9 億", fmt="#,##0.0")
rec("4Q26 營收 vs 高盛（10-19）", "=b_q4", lambda s: "=Q!E4", "高盛（TechNews 2026-10-19）", "", fmt="#,##0.0")
rec("FY2026 EPS vs 高盛（10-19）", "=b_eps26gs", pl("eps", 2026), "高盛（TechNews 2026-10-19）", "")
rec("FY2026 EPS vs 凱基（08-12）", "=b_eps26kgi", pl("eps", 2026), "凱基 2026-08-12", "美銀、大和、花旗 102.7–105.0")
rec("FY2027 營收 vs 凱基（08-12）", "=b_rev27", pl("rev", 2027), "凱基 2026-08-12", "", fmt="#,##0.0")
rec("FY2027 EPS vs 高盛（10-19）", "=b_eps27gs", pl("eps", 2027), "高盛（TechNews 2026-10-19）", "")
rec("FY2027 EPS vs 凱基（08-12）", "=b_eps27kgi", pl("eps", 2027), "凱基 2026-08-12", "")
rec("FY2027 EPS vs 美銀", "=b_eps27bofa", pl("eps", 2027), "時報 2026-08-13", "")
rec("FY2027 EPS vs 大和", "=b_eps27dw", pl("eps", 2027), "時報 2026-08-13", "")
rec("FY2027 EPS vs 花旗（08 月）", "=b_eps27citi", pl("eps", 2027), "時報 2026-08-13", "花旗 09-19 再上修 13%（絕對值未找到）")
rec("FY2028 EPS vs 高盛（10-19）", "=b_eps28gs", pl("eps", 2028), "高盛（TechNews 2026-10-19）", "美銀 178.74、大和 183.06、花旗 185.53（08 月）")
rec("FY2027 毛利率 vs 花旗（百分點）", "=b_gm27", pl("gm", 2027), "經濟日報 2026-09-19", "高盛 33.4%、凱基 33.8%", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "B. 液冷結構（機制檢查）").font = BOLD; rr_ += 1
rec("FY2027 液冷 ÷ 營收 vs 凱基（超過 50%）", "=b_lc27", pl("lcsh", 2027), "凱基 2026-08-12", "", fmt="0.0%", pt=True)
rec("CY2026 液冷滲透率（GW）vs TrendForce（顆數）", "=tf_pen26", lambda s: f"=PL!D{PB[s]['lpen']}", "TrendForce（ETtoday 2026-08-17）", "口徑不同：GW 計高於顆數計", fmt="0.0%", pt=True)
rec("CY2027 液冷滲透率（GW）vs TrendForce（顆數）", "=tf_pen27", lambda s: f"=PL!E{PB[s]['lpen']}", "同上", "", fmt="0.0%", pt=True)
rec("FY2026 奇鋐有效份額（本模型市場口徑）", None, lambda s: f"=PL!D{PB[s]['sh0']}", "—", "份額第二來源：未找到", fmt="0.0%")
rec("CY2027（台達電液冷＋奇鋐液冷）÷ 本模型液冷市場", None, lambda s: f"=(dlc_27+PL!E{PB[s]['lc']}/fx)/PL!E{PB[s]['lmk']}", "台達電模型 Units（唯讀）", "超過 100%：台達電液冷多在機架外（CDU、液對氣），不在本市場口徑內；但也提示本模型的液冷市場口徑可能偏窄、奇鋐份額可能偏高", fmt="0.0%")
rec("CY2026（台達電液冷＋奇鋐液冷）÷ 本模型液冷市場", None, lambda s: f"=(dlc_26+PL!D{PB[s]['lc']}/fx)/PL!D{PB[s]['lmk']}", "同上", "", fmt="0.0%")
rec("FY2027 每 GW 奇鋐 AI 營收（US$M）", None, pl("pgw", 2027), "—", "", fmt="#,##0")
rec("FY2027 AI ÷ 營收", None, pl("aish", 2027), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("其他營收毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "液冷毛利率＝此值＋d_lc", fmt="0.0%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "液冷占比上升使毛利率上升", fmt="0.0%", pt=True)
rec("4Q26 EPS（模型）", None, lambda s: "=Q!E15", "—", "凱基 08-05 估 32.82（以 4Q 營收 626 億為基礎）", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代「AI IT 資本支出 ÷ 每 GW IT 資本支出」"); RC.cell(rr_, 10, "預留列：本版無數值"); RECON_END = rr_; rr_ += 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文經媒體轉述：法說簡報原檔未逐份抓取；券商與研調皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司（經優分析）：4Q25 法說 2026-03-11（FY2025 損益、資本支出、產能、GPU／ASIC 組合）", UA4Q),
        ("公司（經時報）：1Q26 法說 2026-05-14（1Q26 損益、產品與應用比重）", WR1Q),
        ("公司（經優分析）：2Q26 法說 2026-08-12（1H26 損益、伺服器應用比重、3Q／4Q 展望、液冷產能）", UA2Q),
        ("公司（經時報）：2Q26 法說與外資 EPS 2026-08-13（7 月營收、美銀、大和、花旗、中信、群益）", WR2Q),
        ("公司（經 TechNews）：9 月營收 2026-10-07", TN09),
        ("資料商：HiStock 各季利潤比率（稅前淨利率）、股價、股數", HIS),
        ("券商：凱基 2026-08-12（2Q26 實績、估計、產品比重）", KGI08),
        ("券商：凱基 2026-08-05（2025 產品比重、伺服器比重展望）", KGI05),
        ("券商：高盛（經 TechNews 2026-10-19；EPS 2026–28、4Q26 月營收）", TN1019),
        ("券商：澳系（經 TechNews 2026-04-29；液冷占 2026 營收 30–35%）", TN0429),
        ("券商：花旗（經濟日報 2026-09-19；2027 毛利率 34.9%）", "https://money.udn.com/money/amp/story/5607/9764013"),
        ("研調：TrendForce AI 晶片液冷滲透率（經 ETtoday 2026-08-17）", ETTF),
        ("研調：Morgan Stanley NVL72／NVL144 冷卻 BOM（經 Tom's Hardware 2025-11-05；橋接表引用）", "https://finance.yahoo.com/news/cooling-system-single-nvidia-blackwell-154500100.html"),
        ("資料商：Goodinfo 歷年年底收盤與 EPS", "https://goodinfo.tw/tw/StockBzPerformance.asp?STOCK_ID=3017"),
        ("資料商：stockanalysis 年底 P/E（2021–2025）", "https://stockanalysis.com/quote/tpe/3017/financials/ratios/"),
        ("資料商：鉅亨（10-22 收盤）、財報狗（近 12 月 P/E）", "https://www.cnyes.com/twstock/3017"),
        ("市場層：markets/network、markets/chip_bridge、tsmc 成品；台達電模型 Units（Recon 用）（本 repo，唯讀）", ""),
        ("待補（人工）：3Q26 財報與法說（預計 11 月中）；法說簡報原檔；液冷與 AI 的季度金額；奇鋐液冷份額；分歧管單價；2016–2020 年 P/E 第二來源；前瞻 P/E 十年歷史（Bloomberg）；券商共識（付費）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("奇鋐（3017）收支與估值模型 v0.1.1（AI 半導體 Project：電源與散熱次產業；樣板＝台達電）　2026-10-25", True),
 ("版本：v0.1.1（2026-10-25：Val 新增第 0 節「目標價與上檔空間」、In 新增情境權重 tp_wA／tp_wB／tp_wC；其餘數值與 v0.1 相同）。v0.1（初版）：每架液冷內容取自橋接表（Morgan Stanley BOM）；液冷滲透率、每架內容調整、奇鋐份額為本模型 Assumed，在 Units 頁建立。台達電液冷線只在 Recon 對帳。", False),
 ("模型命題：奇鋐的 AI 營收＝部署 GW（AI IT 資本支出 ÷ 每 GW IT 資本支出）×（液冷滲透率 × 每 GW 液冷內容 × 奇鋐份額＋每 GW AI 氣冷與機構）。VR200 每 GW 機架數較少，橋接表的每 GW 液冷內容下降；成長要靠液冷滲透率、Rubin 全液冷與 ASIC 液冷提高每架內容，同時守住冷板份額。", False),
 ("核心問題：液冷滲透率與每架液冷內容的上升，能不能抵銷 VR200 每 GW 機架數下降與份額流失，讓 FY2027–28 EPS 維持券商預期的高成長；現價約 20 倍 FY2027 EPS 隱含的路徑合不合理。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029（三年）。金額：十億新台幣；市場以美元計、以匯率換算。口徑：歸屬母公司淨利。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("液冷營收：markets/network 的 AI IT 資本支出 × 情境顆數係數 ÷ 每 GW IT 資本支出（Tokenomics，經橋接表）＝部署 GW；× 液冷滲透率（Assumed）× 每 GW 液冷內容（橋接表每架機架內液冷＋分歧管 × 每 GW 機架數 × 內容調整）× 匯率 × 奇鋐有效份額（FY2026 以券商液冷占比校準）× 份額相對值（Assumed）。", False),
 ("AI 氣冷與機構：FY2026 以伺服器與網通比重（公司）× AI 比例（Assumed）校準每 GW 營收 × 相對值（氣冷轉液冷，Assumed）。一般伺服器與網通、非伺服器（3C、系統、富世達）：FY2026 拆分 × 成長率（Assumed）。毛利率：液冷比其他營收高 d_lc（Assumed），其他毛利率為 FY2026 殘差（Derived）。營業費用：彈性 phi（Derived）。", False),
 ("校準：FY2026 ＝ 1H26 實績＋3Q26 營收實績（月營收）＋假設（4Q 季增、毛利率、營業費用率）。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：現價 ÷ FY2027、FY2028 EPS；十年 P/E 規則值（Goodinfo 年底收盤 ÷ EPS；stockanalysis 核對）× FY2027 EPS；循環位置以本模型營收最高年判定（奇鋐不套設備股見頂年）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文經媒體轉述；券商多經媒體轉述。", False),
 ("", False),
 ("已知限制：(1) 2026 液冷占比只有券商說法（30–35%），公司未揭露；(2) 液冷滲透率、每架內容調整、份額相對值為 Assumed，是 FY2028 EPS 最敏感的輸入；(3) 部署 GW 以 VR200 每 GW IT 資本支出換算，ASIC 機架的每 GW 資本支出較低，GW 可能低估；(4) 十年 P/E 的 2016–2022 年為 AI 前的低倍數（約 10–15x），規則值可能低估 AI 期的合理倍數；(5) 稅前利益與 FY2025 各季營收為推導；(6) 3Q26 財報（預計 11 月中）後更新錨點；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

V2["pe10"] = (list(zip(range(2016, 2026), PE_GI)), "Goodinfo 年底收盤 ÷ 當年 EPS；同規則值")
V2["old_row"] = TP0
# ================= v0.2（2026-10-25）：PL 延伸到 FY2032、Val 第 0 節目標價（DCF 與 Comparable 加權）、Sens 目標價參數 =================
# 全部放在各頁既有內容之後，不位移既有儲存格（FY2025–FY2029 數值與 v0.1.1 相同）
FADEFILL = PatternFill("solid", fgColor="E7E6E6")
XY = (2025, 2026, 2027, 2028, 2029, 2030, 2031, 2032); XC = {y: "CDEFGHIJ"[i] for i, y in enumerate(XY)}; FADE = (2030, 2031, 2032)
K2 = V2["keys"]
for rw in range(4, IN.max_row + 1):
    if IN.cell(rw, 1).value in ("tp_wA", "tp_wB", "tp_wC", "tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "nim10_f"):
        IN.cell(rw, 5).fill = KEY
# --- PL 延伸表（三情境到 FY2032） ---
xr = PL.max_row + 3; XT = xr
PL.cell(xr, 1, "延伸：三情境到 FY2032（v0.2 新增）。FY2025–FY2029 引用上方推導區塊；FY2030 起為收斂法（灰底），非驅動因子推導：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 線性收斂到收斂目標，FY2032 達到").font = BOLD
xr += 1; XR = {}; XH = None
for s in SC:
    E = PB[s]; R = {}
    PL.cell(xr, 1, f"延伸｜{SCN_T[s]}").font = BOLD; xr += 1
    hdr(PL, xr, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    XH = XH or xr; xr += 1
    for key, lab, unit in (("rev", "營收（延伸）", V2["unit"]), ("g", "  營收年增", "%"), ("ni", "淨利（延伸）", V2["unit"]), ("nim", "  淨利率", "%"),
                           ("sh", "股數（延伸）", V2["shunit"]), ("eps", "EPS（延伸）", V2["cur"]), ("epsg", "  EPS 年增", "%"),
                           ("tg", "收斂目標：FY2032 營收年增", "%"), ("tm", "收斂目標：FY2032 淨利率", "%")):
        PL.cell(xr, 1, lab); PL.cell(xr, 2, unit); R[key] = xr; xr += 1
    G = lambda k: f"$G${R[k]}"
    for y in XY:
        c = XC[y]; pc = XC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={c}{E[K2['rev']]}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            PL[f"{c}{R['ni']}"] = f"={c}{E[K2['ni']]}"
            if y > 2025 or K2.get("sh25", True): PL[f"{c}{R['sh']}"] = f"={c}{E[K2['sh']]}"
            PL[f"{c}{R['eps']}"] = f"={c}{E[K2['eps']]}"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"={G('g')}+($C${R['tg']}-{G('g')})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"={G('nim')}+($C${R['tm']}-{G('nim')})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['sh']}"] = f"={pc}{R['sh']}*(1+{V2['shg']})"
            PL[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh']}"
        if y <= 2029: PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
        if y > 2025: PL[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("rev", "#,##0.0"), ("g", "0.0%"), ("ni", "#,##0.0"), ("nim", "0.0%"), ("sh", "0.0000"), ("eps", V2["epsfmt"]), ("epsg", "0.0%")):
            PL[f"{c}{R[key]}"].number_format = fm
            if y in FADE: PL[f"{c}{R[key]}"].fill = FADEFILL
    PL[f"C{R['tg']}"] = f"=fade_g{s}"; PL[f"C{R['tg']}"].number_format = "0.0%"
    PL[f"C{R['tm']}"] = {"A": f"=G{R['nim']}", "B": f"=G{R['nim']}*(1+nim10_f)/2", "C": f"=G{R['nim']}*nim10_f"}[s]; PL[f"C{R['tm']}"].number_format = "0.0%"
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導（上游市場層只到 2029）")
    PL.cell(R["tg"], 11, {"A": "In 頁 fade_gA＝tp_g＋2pt", "B": "In 頁 fade_gB＝tp_g", "C": "In 頁 fade_gC＝tp_g−2pt"}[s])
    PL.cell(R["tm"], 11, {"A": "＝本情境 FY2029 淨利率", "B": "＝FY2029 淨利率與十年平均替代值（FY2029 × nim10_f）的中點", "C": "＝十年平均替代值（FY2029 × nim10_f；十年平均未找到）"}[s])
    PL.cell(R["sh"], 11, f"FY2030 起股數年變動＝{V2['shg']}")
    XR[s] = R; xr += 1
PL.column_dimensions["I"].width = 14; PL.column_dimensions["J"].width = 14; PL.column_dimensions["K"].width = 60
# --- Val 第 0 節：目標價 ---
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
t0 = V.max_row + 3; T0 = t0
V.cell(t0, 1, "0. 目標價（DCF 與 Comparable 加權；v0.2 新增，放在本頁既有內容之後以免位移既有儲存格；三情境欄 B／C／D＝A 順利／B 平平／C 不如預期）").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["h1"] = t0 + 2
r_ = t0 + 3
V.cell(r_, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{r_}"] = "PL 頁延伸表；FY2030 起為收斂法"; T["eps_h"] = r_; r_ += 1
EPSR = {}
for y in YRS6:
    tset(r_, f"eps{y}", f"  FY{y}", [f"=PL!{XC[y]}{XR[s]['eps']}" for s in SC], V2["epsfmt"]); EPSR[y] = r_; r_ += 1
V.cell(r_, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = r_; r_ += 1
tset(r_, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）"); RV = r_; r_ += 1
CFR = {}
for y in YRS6:
    tset(r_, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], V2["epsfmt"]); CFR[y] = r_; r_ += 1
def dcf_formula(c, r="tp_r", g="tp_g", reinv=None):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    reinv = reinv or f"Val!$B${RV}"
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
c1, c6 = CFR[2027], CFR[2032]
tset(r_, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], V2["pxfmt"],
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027"); r_ += 1
tset(r_, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{T['dcfv']}*(1+tp_r)" for c in "BCD"], V2["pxfmt"]); r_ += 1
tset(r_, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{T['dcf12']}+tp_wB*C{T['dcf12']}+tp_wC*D{T['dcf12']}", None, None], V2["pxfmt"], "Σ tp_wA／wB／wC × 情境價值 ×(1＋tp_r)；權重見 In 頁"); r_ += 2
V.cell(r_, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; r_ += 1
hdr(V, r_, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["ph"] = r_; r_ += 1
P1 = r_
for nm, tk, pe, dt in V2["peers"]:
    V.cell(r_, 1, nm); V.cell(r_, 2, tk); cc = V.cell(r_, 3, pe); cc.font = BLUE; cc.number_format = "0.0"
    V.cell(r_, 4, dt); V.cell(r_, 5, "stockanalysis.com 個股頁 Forward PE（分析師共識；價格日期如左）"); V.cell(r_, 6, "Interested-party（賣方共識）"); r_ += 1
P2 = r_ - 1; T["p1"], T["p2"] = P1, P2
tset(r_, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{P1}:C{P2})", None, None], "0.0", V2["peer_note"]); r_ += 1
w1, w2 = V2["fwd_w"]
tset(r_, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*({w1}*{c}{EPSR[2027]}+{w2}*{c}{EPSR[2028]})" for s, c in zip(SC, "BCD")), None, None], V2["epsfmt"], V2["fwd_note"]); r_ += 1
tset(r_, "comptp", "Comparable 目標價", [f"=B{T['pmed']}*B{T['feps']}", None, None], V2["pxfmt"], "同業中位數 × 上列 EPS"); r_ += 2
V.cell(r_, 1, "0.3 結論").font = BOLD; r_ += 1
hdr(V, r_, ["項目", "數值", "", "", "說明"]); T["h3"] = r_; r_ += 1
tset(r_, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], V2["pxfmt"], "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁"); r_ += 1
tset(r_, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], V2["pxfmt"], "引用 0.2"); r_ += 1
tset(r_, "cons", "市場共識目標價", ["=cons_tp", None, None], V2["pxfmt"], V2["cons_tp"][1]); r_ += 1
tset(r_, "upw", "上檔空間：加權目標價", [f"=B{T['tpw']}/px-1", None, None], "0.0%", V2["px_note"]); r_ += 1
tset(r_, "upc", "上檔空間：Comparable 目標價", [f"=B{T['comp2']}/px-1", None, None], "0.0%"); r_ += 1
tset(r_, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS 見 In 頁 cons_feps 與 Recon"); r_ += 1
tset(r_, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%"); r_ += 1
V[f"E{T['gmul']}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(r_, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）"); r_ += 1
PEROW = r_ + 3
PE10 = f"$B${PEROW + 1}:$K${PEROW + 1}"
tset(r_, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
     [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E 分布（下方列）的分位；低於最小值記 0%、高於最大值記 100%；只作檢查，不進目標價"); r_ += 1
V.cell(r_, 1, "目標價方法"); V[f"B{r_}"] = V2["method"]; T["m"] = r_; r_ += 1
V.cell(PEROW, 1, "十年年底 P/E（歷史分位檢查用；" + V2["pe10"][1] + "）").font = BOLD
for j, (yy, vv) in enumerate(V2["pe10"][0]):
    V.cell(PEROW, 2 + j, str(yy))
    cc = V.cell(PEROW + 1, 2 + j, vv); cc.font = BLUE; cc.number_format = "0.0"
V.cell(PEROW + 1, 1, "年底 P/E"); T["pe10h"] = PEROW; T["pe10"] = PEROW + 1
for k_ in ("dcftp", "comptp", "tpw"):
    V[f"B{T[k_]}"].fill = KEY
V.cell(T["tpw"], 1).font = BOLD
defname("tp_final", f"Val!$B${T['tpw']}"); defname("tp_dcf", f"Val!$B${T['dcftp']}"); defname("tp_comp", f"Val!$B${T['comptp']}")
V["A3"] = f"目標價：見第 0 節（本頁第 {T0} 列；v0.2 新增，DCF 與 Comparable 加權）。第 {V2['old_row']} 列「舊法：規則值情境加權（參考）」不再是目標價，只作參考"
for c in "IJK": V.column_dimensions[c].width = 10
# --- Sens：目標價參數 ---
TPS0 = SE.max_row + 3
SE.cell(TPS0, 1, "目標價參數（v0.2）：一次動一個，對加權目標價的影響（只改 DCF 參數；PL 收斂目標 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = V2["pxfmt"]
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "tp_r 10%、tp_g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "tp_r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "tp_r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "tp_g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "tp_g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.4", "tp_w_dcf 0.4", dict(w="0.4")), ("DCF 權重 0.8", "tp_w_dcf 0.8", dict(w="0.8"))]
for k_, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k_, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
# --- Recon：目標價對帳 ---
rq = RC.max_row + 2
RC.cell(rq, 1, "E. 目標價對帳（v0.2；模型欄為情境加權值，非情境 A）").font = BOLD; rq += 1
RECV2 = rq
for lab, ref, mdl, src in (("12 個月後的前瞻 EPS（情境加權）vs 共識前瞻 EPS", "=cons_feps", f"=Val!B{T['feps']}", V2["cons_feps"][1]),
                           ("目標價（DCF 與 Comparable 加權）vs 市場共識目標價", "=cons_tp", f"=Val!B{T['tpw']}", V2["cons_tp"][1]),
                           ("Comparable 目標價 vs 市場共識目標價", "=cons_tp", f"=Val!B{T['comptp']}", "同上")):
    RC.cell(rq, 1, lab); RC.cell(rq, 2, ref).number_format = V2["epsfmt"]; RC.cell(rq, 3, mdl).number_format = V2["epsfmt"]
    RC.cell(rq, 4, f"=C{rq}/B{rq}-1").number_format = "0.0%"; RC.cell(rq, 9, src); RC.cell(rq, 10, "共識只作對帳，不作校準"); rq += 1
# --- README ---
for rw in range(1, RD.max_row + 1):
    v_ = RD.cell(rw, 1).value
    if isinstance(v_, str) and rw == 1: RD.cell(rw, 1, v_.replace("v0.1.1", "v0.2.1"))
    elif isinstance(v_, str) and v_.startswith("估值："): RD.cell(rw, 1, V2["readme_val"] + "（v0.1.1 原估值說明，保留作參考：" + v_[3:] + "）")
    elif isinstance(v_, str) and v_.startswith("版本："): RD.cell(rw, 1, "版本：v0.2.1（2026-10-25：液冷滲透率改依情境字母（奇鋐），數值變動：FY2027 起 A＝高值、B＝基準、C＝低值（與雙鴻模型同值，Andy 2026-10-25 13:13）；補上檔空間級距、投資建議、重點觀察）。v0.2（2026-10-25：目標價改為 DCF 與 Comparable 加權、三情境延伸到 FY2032；FY2025–FY2029 數值不變）。前版：" + v_[3:])
rn = RD.max_row + 1
for t_ in V2["readme_add"]:
    RD.cell(rn, 1, t_); rn += 1
if os.environ.get("V2_ROWMAP"):
    import json as _json
    _json.dump({"XT": XT, "XH": XH, "XR": XR, "T": T, "EPSR": EPSR, "CFR": CFR, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "RECV2": RECV2},
               open(os.environ["V2_ROWMAP"], "w"), ensure_ascii=False, indent=1)

# ================= v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（工作單；Andy 2026-10-25 14:00、14:17） =================
# 全部放在 Val、README 現有最後一列之後，不位移既有儲存格
OBS = [
       "目標價 NT$3,966（DCF 3,002、Comparable 4,929 各半），比現價 3,145（2026-10-22）高 26.1%，級距 25%–50%；Comparable 單獨為 56.7%。主因是 EPS：本模型 12 個月後前瞻 EPS 179.1 比共識前瞻 EPS 132.6 高 35.1%（兩者期間差約一年；FY2027 B 163.2 比高盛 154.8 高 5.4%），同業（雙鴻、健策、Vertiv、台達電）前瞻 P/E 中位數 27.5 倍略高於現價隱含 23.7 倍；DCF 只有前瞻 EPS 的 16.8 倍、比現價低 4.5%，上檔空間全由 Comparable 拉動。",
       "v0.2.1 起液冷滲透率依情境字母（A 高、B 基準、C 低）：FY2027 EPS A／B／C＝174.7／163.2／145.8（原三情境同為 163.2），FY2028＝224.8／177.5／130.4（原 209.7／177.5／143.4）；加權目標價由 3,956 升到 3,966，上檔空間由 25.8% 升到 26.1%，離買進與中立的界線 +25% 很近。",
       "每架液冷內容調整是最敏感的變數：調整 1.0（只有橋接表的機架內液冷）時 FY2028 EPS −18.2%，2027 1.6、2028 1.8 時 +18.3%。10 月營收（約 11-10）與 3Q26 財報與法說（預計 11 月中）可檢驗 2026 年液冷占比（本模型 32%，券商 30–35%）。"
      ]   # 重點觀察（與 README.md 相同文字；數字取自重算後的 xlsx）
_U = f"B{T['upw']}"
GR = V.max_row + 3
V.cell(GR - 1, 1, "0.4 上檔空間級距與投資建議（v0.2.1 新增；放在本頁最後以免位移既有儲存格；依「上檔空間：加權目標價」第 " + str(T["upw"]) + " 列）").font = BOLD
tset(GR, "grade", "上檔空間級距", [f'=IF({_U}>0.5,"大於 50%",IF({_U}>0.25,"25%–50%",IF({_U}>=0,"0%–25%",IF({_U}>=-0.25,"0%～−25%",IF({_U}>=-0.5,"−25%～−50%","小於 −50%")))))', None, None], "General",
     "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）；負值區間同樣重要")
tset(GR + 1, "rec", "投資建議", [f'=IF({_U}>0.5,"強力買進",IF({_U}>0.25,"買進",IF({_U}>=-0.25,"中立",IF({_U}>=-0.5,"放空","強力放空"))))', None, None], "General",
     "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）")
rn = RD.max_row + 1
RD.cell(rn, 1, f'="投資建議："&Val!B{GR + 1}&"（依上檔空間固定規則）"').font = BOLD
OBS_ROW = rn + 1
for i_, t_ in enumerate(OBS, 1):
    RD.cell(rn + i_, 1, f"重點觀察 {i_}：" + t_)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("AVC_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    IN_ROWS = {}
    rr = 4
    for srow in rows:
        if srow[0] != "sec": IN_ROWS[srow[0]] = rr
        rr += 1
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK_ROW, "SENSN": SENS_N, "SENSNOTE": SENS_NOTE, "VAL": VAL, "SB": SB,
               "IN": IN_ROWS, "VR_NOTE": vr - 1, "BLOCK_END": block_end, "RECON_END": RECON_END, "UN_NOTE": ur},
              open(os.environ["AVC_ROWMAP"], "w"), ensure_ascii=False, indent=1)
