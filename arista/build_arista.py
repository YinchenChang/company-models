# Arista Networks（ANET）收支與估值模型 v0.1（AI 半導體 Project：網通系統端樣板）
# 用法：python3 arista/build_arista.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不做平移。金額：十億美元（US$B）；EPS：美元；口徑：公司 non-GAAP（排除股票薪酬、攤銷等；
# 2026 年起 non-GAAP 稅額改用新方法，公司已重編 2025 年各季，本模型 FY2025 用 2025 年度新聞稿的原報值，見 Q 頁說明）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
wsn = openpyxl.load_workbook(NET_F, data_only=True)
NETR = {}
for r in wsn["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in ("scale-out 交換器", "AI IT 資本支出（加速器 × 每顆 IT capex）", "NVIDIA 網通（本模型歸屬）", "AI 網通市場合計"):
        NETR[lab] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETR) == 4 and all(v is not None for v in NETR["scale-out 交換器"][0]), NETR.keys()
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝Arista 8-K 新聞稿（SEC EDGAR）；公司的 AI 營收目標與全年指引為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
EDG = "https://www.sec.gov/Archives/edgar/data/0001596532/"
PR_Q226 = EDG + "000159653226000174/ex991q226-earningsrelease.htm"
PR_Q425 = "https://www.businesswire.com/news/home/20260212962784/en/Arista-Networks-Inc.-Reports-Fourth-Quarter-and-Year-End-2025-Financial-Results"
CALL = "法說轉述：BigGo Finance 摘要（2026-08-04；二手）"

sec("A. AI 後端 scale-out 交換器市場（讀 markets/network 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mkt_{y}", f"AI 後端 scale-out 交換器市場 20{y}", "US$B", None, NETR["scale-out 交換器"][0][j], None, "Derived", NETR["scale-out 交換器"][1],
        "加速器顆數 × 每顆 scale-out 交換器內容（不分乙太網路與 InfiniBand）；只含 AI 後端" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][0][j], None, "Derived", NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][1], "Recon 用：Arista AI 營收 ÷ AI IT 資本支出" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、Broadcom 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")

sec("B. 乙太網路占比與 Arista 份額（本專案為權威：晶片商與網通廠份額，規則第 13 條）")
inp("eth_26", "乙太網路 ÷ AI 後端交換器（金額）2026", "%", 0.70, 0.75, 0.80, "Assumed", "Dell'Oro 2026-03-10：2025 年乙太網路占 AI 叢集交換器三分之二以上、為 InfiniBand 兩倍以上（Interested-party）", "IDC 2Q26：NVIDIA 資料中心乙太網路年增 181%（乙太網路取代 InfiniBand 的方向）")
inp("eth_27", "乙太網路 ÷ AI 後端交換器 2027", "%", 0.72, 0.80, 0.88, "Assumed", "同上；2027 年無公開預測：未找到", "外推（Assumed）")
inp("eth_28", "乙太網路 ÷ AI 後端交換器 2028", "%", 0.74, 0.83, 0.90, "Assumed", "同上", "—")
inp("eth_29", "乙太網路 ÷ AI 後端交換器 2029", "%", 0.75, 0.85, 0.92, "Assumed", "同上", "—")
inp("rel_27", "Arista 份額 2027 ÷ 2026（相對值）", "x", 0.80, 0.95, 1.05, "Assumed", "IDC 2Q26：NVIDIA 20.4% 已超過 Arista 18.7%（資料中心乙太網路，全部速率；Interested-party）", "Dell'Oro 2025 AI 後端乙太網路：Celestica（白牌）與 NVIDIA 並列第一、合計約 50%，Arista 第三（Interested-party）")
inp("rel_28", "Arista 份額 2028 ÷ 2026", "x", 0.70, 0.90, 1.05, "Assumed", "同上", "—")
inp("rel_29", "Arista 份額 2029 ÷ 2026", "x", 0.65, 0.88, 1.05, "Assumed", "同上", "—")

sec("C. 非 AI 營收（雲端前端、企業、園區、服務）")
inp("g_nai27", "非 AI 營收年增 2027", "%", 0.03, 0.10, 0.15, "Assumed", "2026 錨點非 AI 約 9.1（營收指引 12.6 − AI 目標 3.5）；2025 約 7.4（9.0 − AI 約 1.63）→ 2026 約 +23%（Derived，含記憶體缺貨下的提前拉貨可能）", "園區目標 2026 至少 1.25（公司，Interested-party）")
inp("g_nai28", "非 AI 營收年增 2028", "%", 0.02, 0.08, 0.12, "Assumed", "同上", "—")
inp("g_nai29", "非 AI 營收年增 2029", "%", 0.02, 0.06, 0.10, "Assumed", "同上", "—")
inp("ai25", "2025 年 AI 營收（參考）", "US$B", 1.5, 1.63, 1.8, "Derived", "Next Platform 2026-02-13：以 2026 目標倒推約 1.63（公司稱超過 1.5 目標）", "公司：2025 年 AI 目標 1.5 已超越（法說）")

sec("D. 毛利率")
inp("d_ai", "AI 營收毛利率低於非 AI 的幅度", "pt", 0.03, 0.06, 0.10, "Assumed", "公司：毛利率下滑來自雲端與 AI 大客戶占比（2Q26 毛利率較去年同期 −2.2 個百分點；法說，二手）", "產品別毛利率未揭露：未找到")

sec("E. FY2026 錨點（Q1–Q2 實績＋Q3 指引＋全年指引）")
inp("fy_rev", "FY2026 營收指引", "US$B", None, 12.6, None, "Interested-party", CALL + "：約 $12.6B（8 月由 11.5 上修）", "StockAnalysis 共識 12.68（27 位）")
inp("fy_gm", "FY2026 non-GAAP 毛利率指引（中點）", "%", 0.62, 0.63, 0.64, "Interested-party", CALL + "：62–64%", "sec-api 彙整同值")
inp("fy_oim", "FY2026 non-GAAP 營業利益率指引（中點）", "%", 0.48, 0.485, 0.49, "Interested-party", CALL + "：48–49%", "sec-api 彙整同值")
inp("q3_rev", "3Q26 營收指引", "US$B", None, 3.3, None, "Interested-party", PR_Q226 + "：約 $3.3B", "BigGo（二手）同值")
inp("q3_gm", "3Q26 non-GAAP 毛利率指引", "%", 0.62, 0.63, 0.64, "Interested-party", CALL + "：約 63%", "第二來源：未找到")
inp("q3_oim", "3Q26 non-GAAP 營業利益率指引（中點）", "%", 0.48, 0.485, 0.49, "Interested-party", PR_Q226 + "：48–49%", "—")
inp("q_oth", "3Q、4Q26 non-GAAP 其他收益（每季）", "US$B", 0.12, 0.13, 0.14, "Derived", "1Q26 0.1136、2Q26 0.1263（公司，GAAP 其他收益）", "現金與有價證券 13.3（2026-06，公司）")
inp("q_tx", "3Q、4Q26 non-GAAP 有效稅率", "%", 0.20, 0.21, 0.22, "Derived", "2Q26：1 − 1301.7 ÷（1513.7＋126.3）＝20.6%；1Q26 21.3%（推導）", "3Q26 EPS 指引 1.06–1.08 隱含約 21%（推導）")
inp("q_sh", "3Q、4Q26 稀釋股數", "十億股", None, 1.277, None, "Derived", "2Q26 1.2760（公司）", "—")
inp("ai26", "FY2026 AI 營收（公司目標「至少」）", "US$B", 3.5, 3.5, 4.0, "Interested-party", CALL + "：AI 至少 $3.5B（2 月 3.25、8 月 3.5）；其中 scale-across 約 1.2", "sec-api 彙整：至少 $3.50B（2026-08-04 法說）")

sec("F. 營業費用、其他收益、稅、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.40, 0.55, 0.80, "Derived", "FY2026 錨點營業費用年增 ÷ 營收年增（Q 頁）", "低／高 Assumed")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 ±2 個百分點")
inp("oth_f", "non-GAAP 其他收益（每年）", "US$B", 0.40, 0.55, 0.70, "Assumed", "FY2026 錨點約 0.50（Q 頁）；現金 13.3 × 約 4%", "—")
inp("tx_f", "non-GAAP 有效稅率", "%", 0.19, 0.21, 0.22, "Derived", "FY2025 19.5%、1H26 約 21%（推導）", "—")
inp("sh_g", "稀釋股數年變動", "%", -0.01, 0.0, 0.005, "Derived", "FY2025 1.2757 → 2Q26 1.2760（公司）：買回大致抵銷股票薪酬", "低／高 Assumed")
inp("mf", "AI 市場係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 scale-out 交換器 +90%；公司稱供給吃緊是「兩年的問題」、2028 年才解除（法說，二手）", "±20%")

sec("G. 共識、市場資料（只作 Recon 對帳，不作校準）")
inp("cons_rev26", "共識 FY2026 營收", "US$B", None, 12.68, None, "Interested-party", "StockAnalysis（27 位，2026-10-23）", "MarketBeat：FY2026 EPS 共識 3.79（口徑可能為 GAAP 或舊值）")
inp("cons_eps26", "共識 FY2026 non-GAAP EPS", "US$", None, 4.12, None, "Interested-party", "StockAnalysis（區間 4.05–4.22）", "—")
inp("erste_eps27", "Erste Group FY2027 EPS（單一分析師）", "US$", None, 4.79, None, "Interested-party", "MarketBeat 2026-10-23（10-21 由 4.74 上調）", "2027 共識（StockAnalysis 付費）：未找到")
inp("dell_top2", "2025 AI 後端乙太網路：Celestica＋NVIDIA 合計份額", "%", None, 0.50, None, "Interested-party", "Dell'Oro 2026-03-10：兩家並列第一、合計約 50%；Arista 第三", "—")
inp("idc_anet", "2Q26 資料中心乙太網路交換器：Arista 份額（全部速率）", "%", None, 0.187, None, "Interested-party", "IDC 2026-09-11", "NVIDIA 20.4%、Cisco 18.2%")
inp("pe_fair_andy", "Andy 先驗：合理 P/E", "x", None, 25, None, "Assumed", "先驗_Andy_Arista_20261024 A9b（Drive；不在本 repo）", "只作對照")

sec("H. 價格與估值")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 172.37, None, "Interested-party", "StockAnalysis（2026-10-23 收盤）", "MarketBeat：10-23 開盤 170.57")
inp("pe_ttm_gaap", "近 12 月 GAAP P/E", "x", None, 54.55, None, "Interested-party", "StockAnalysis", "MarketBeat 53.81")
inp("pe_fwd_sa", "前瞻 P/E（StockAnalysis）", "x", None, 37.06, None, "Interested-party", "StockAnalysis（口徑：未來 12 個月 non-GAAP，頁面未明示）", "第二來源：未找到")
inp("g2n", "GAAP EPS ÷ non-GAAP EPS（規則值換算用）", "x", None, "=1.75/1.89", None, "Derived", PR_Q226 + "：1H26 GAAP 稀釋 EPS 1.75、non-GAAP 1.89", "FY2025：2.75 ÷ 2.98＝0.923（公司）")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（十億美元；EPS 為美元；股數為十億股）。口徑：公司 non-GAAP。藍字＝輸入；黑字＝公式。實績取自 Arista 8-K 新聞稿（SEC EDGAR、Business Wire）"
Q["A2"] = "FY2026 錨點＝1Q26、2Q26 實績＋3Q26 指引＋4Q26（＝全年指引 − 前三季；毛利率、營業利益率同法）。1Q26＝上半年合計 − 2Q26。FY2025 non-GAAP 淨利與 EPS 為 2025 年度新聞稿原報值；公司 2026 年起改 non-GAAP 稅額方法並重編 2025 年各季，重編後的全年值：未找到"
hdr(Q, 3, ["項目", "FY2025A", "1Q26", "2Q26", "3Q26（指引）", "4Q26（全年指引 − 前三季）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：" + PR_Q425 + "；1Q26、2Q26：" + PR_Q226 + "（2Q26 與上半年合計）"
S2 = "1Q26＝上半年 − 2Q26；交叉：sec-api 彙整的 GAAP、non-GAAP 毛利率與營業利益率同值"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
qrow(4, "營收", {"B": 9.0057, "C": "=5.7447-D4", "D": 3.0357, "E": "=q3_rev", "F": "=fy_rev-C4-D4-E4"}, tag="Verified／Interested-party（3Q、4Q 指引）")
qrow(5, "non-GAAP 毛利", {"B": 5.8154, "C": "=3.6155-D5", "D": 1.9247, "E": "=E4*q3_gm", "F": "=fy_rev*fy_gm-C5-D5-E5"}, tag="Verified／Interested-party")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2025 64.6%；FY2026 指引 62–64%")
qrow(7, "non-GAAP 營業費用", {c: f"={c}5-{c}8" for c in "BCDEF"}, tag="Derived（毛利 − 營業利益）")
qrow(8, "non-GAAP 營業利益", {"B": 4.3370, "C": "=2.8077-D8", "D": 1.5137, "E": "=E4*q3_oim", "F": "=fy_rev*fy_oim-C8-D8-E8"}, tag="Verified／Interested-party")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(10, "其他收益（GAAP，non-GAAP 沿用）", {"B": 0.3936, "C": "=0.2399-D10", "D": 0.1263, "E": "=q_oth", "F": "=q_oth"}, tag="Verified／Derived（3Q、4Q）")
qrow(11, "稅前利益", {c: f"={c}8+{c}10" for c in "BCDEF"}, tag="Derived")
qrow(12, "non-GAAP 淨利", {"B": 3.8060, "C": "=2.4099-D12", "D": 1.3017, "E": "=E11*(1-q_tx)", "F": "=F11*(1-q_tx)"}, tag="Verified／Derived（3Q、4Q）")
qrow(13, "有效稅率（推導）", {c: f"=1-{c}12/{c}11" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="non-GAAP 稅額未直接揭露：1 − 淨利 ÷ 稅前")
qrow(14, "稀釋股數（十億股）", {"B": 1.2757, "C": "=(1.2749*2-D14)", "D": 1.2760, "E": "=q_sh", "F": "=q_sh"}, ann="=AVERAGE(C14:F14)", tag="Verified／Derived", s2="1Q26＝上半年加權 × 2 − 2Q26（近似）")
qrow(15, "non-GAAP EPS", {c: f"={c}12/{c}14" for c in "BCDEF"}, fmt="0.000", ann="=G12/G14", tag="Derived", s1="淨利 ÷ 稀釋股數", s2="—")
qrow(16, "  公司揭露 non-GAAP EPS", {"B": 2.98, "D": 1.02}, fmt="0.00", ann="", tag="Verified", s2="1Q26：1.89 − 1.02＝0.87；3Q26 指引 1.06–1.08")
qrow(17, "AI 營收", {"B": "=ai25"}, ann="=ai26", tag="Derived／Interested-party", s1="FY2025：In ai25；FY2026：公司目標「至少 3.5」", s2="各季 AI 營收未揭露：未找到")
qrow(18, "非 AI 營收（＝營收 − AI）", {"B": "=B4-B17"}, ann="=G4-G17", tag="Derived", s1="—", s2="—")
qrow(19, "AI ÷ 營收", {"B": "=B17/B4"}, fmt="0.0%", ann="=G17/G4", tag="Derived", s1="—", s2="Andy 先驗 A2：2026 年 50%")
Q.cell(21, 1, "FY2026 錨點與共識的距離（參考；校準未使用共識）").font = BOLD
Q.cell(22, 1, "FY2026 營收 ÷ 共識 − 1"); Q["B22"] = "=G4/cons_rev26-1"
Q.cell(23, 1, "FY2026 EPS ÷ 共識 − 1"); Q["B23"] = "=G15/cons_eps26-1"
Q.cell(24, 1, "FY2026 營收年增"); Q["B24"] = "=G4/B4-1"
Q.cell(25, 1, "FY2026 營業費用年增"); Q["B25"] = "=G7/B7-1"
Q.cell(26, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B26"] = "=B25/B24"
Q.cell(27, 1, "3Q26 EPS（模型）vs 公司指引中點 1.07"); Q["B27"] = "=E15/1.07-1"
for rr in range(22, 28): Q[f"B{rr}"].number_format = "0.0%" if rr != 26 else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("oth", 10), ("ni", 12), ("sh", 14), ("eps", 15), ("ai", 17), ("nai", 18)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 52
for c in "BCDEFG": Q.column_dimensions[c].width = 15
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 90; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", rel="rel", eth="eth", mf="1", gn="g_nai", d="d_ai", opxa="0")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 營收（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029）")
    row("mkt", "AI 後端 scale-out 交換器市場 × 情境顆數係數", "US$B", {c_(y): f"=mkt_{y % 100}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("eth", "乙太網路占比", "%", {c_(y): f"={P['eth']}_{y % 100}" for y in BU}, fmt="0.0%")
    row("emk", "AI 後端乙太網路交換器市場", "US$B", {c_(y): f"={c_(y)}{R['mkt']}*{c_(y)}{R['eth']}" for y in BU})
    row("sh0", "Arista 有效份額（FY2026 校準）＝ AI 目標 ÷ 乙太網路市場", "%", {"D": f"=a26_ai/D{R['emk']}"}, fmt="0.0%",
        note="Derived：分子含前端與 scale-across，所以高於純後端份額；FY2027 起乘相對值")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "Arista 有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("ai", "AI 營收", "US$B", {"C": "=a25_ai", **{c_(y): f"={c_(y)}{R['emk']}*{c_(y)}{R['sh']}" for y in BU}})
    sec_("B. 非 AI 與合計")
    row("nai", "非 AI 營收", "US$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']}{y % 100})" for y in BU[1:]}})
    row("rev", "營收", "US$B", {"C": "=a25_rev", **{c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in BU}})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    sec_("C. 損益")
    row("gmn", "非 AI 毛利率 ＝ 錨點毛利率 ＋ 差距 × AI 占比（FY2026 校準）", "%", {"D": f"=a26_gp/a26_rev+{P['d']}*D{R['aish']}"}, fmt="0.0%",
        note="Derived：讓 FY2026 毛利率等於錨點；AI 毛利率＝非 AI − 差距。FY2027 起毛利率隨 AI 占比變動")
    row("gp", "毛利", "US$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"=$D${R['gmn']}*{c_(y)}{R['nai']}+($D${R['gmn']}-{P['d']})*{c_(y)}{R['ai']}" for y in BU[1:]}})
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in BU[1:]}})
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}})
    row("oi", "營業利益", "US$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("oth", "其他收益", "US$B", {"C": "=a25_oth", "D": "=a26_oth", **{c_(y): "=oth_f" for y in BU[1:]}})
    row("etr", "有效稅率", "%", {"C": "=1-a25_ni/(a25_oi+a25_oth)", "D": "=1-a26_ni/(a26_oi+a26_oth)", **{c_(y): "=tx_f" for y in BU[1:]}}, fmt="0.0%")
    row("ni", "淨利（non-GAAP）", "US$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}+{c_(y)}{R['oth']})*(1-{c_(y)}{R['etr']})" for y in BU[1:]}})
    row("sh_", "稀釋股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.000")
    row("eps", "EPS（non-GAAP）", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Arista 損益（US$B；non-GAAP）。FY2025A 實績、FY2026 錨點（1Q–2Q 實績＋3Q 指引＋全年指引），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.00"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 營收", "US$B", "#,##0.00"), ("nai", "  非 AI 營收", "US$B", "#,##0.00"),
                                ("aish", "  AI ÷ 營收", "%", "0.0%"), ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利（non-GAAP）", "US$B", "#,##0.00"),
                                ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收 − 錨點；FY2026 模型 AI − 公司目標（校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：AI 差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['ai']}-a26_ai"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['gmn']}*D{PB['A']['nai']}+(D{PB['A']['gmn']}-d_ai)*D{PB['A']['ai']}-a26_gp"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：份額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — Arista 在 AI 後端乙太網路交換器的份額（本專案權威，規則第 13 條）。市場取自 markets/network（只含 AI 後端 scale-out 交換器，不分協定）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("mkt", "AI 後端 scale-out 交換器市場（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mkt']}", "#,##0.0", "markets/network Mkt 頁"),
        ("eth", "乙太網路占比", "%", lambda y: f"=PL!{FC[y]}{A_['eth']}", "0.0%", "In 頁 B 節（Assumed）"),
        ("emk", "AI 後端乙太網路交換器市場", "US$B", lambda y: f"=PL!{FC[y]}{A_['emk']}", "#,##0.0", ""),
        ("ib", "InfiniBand 等其他協定（＝市場 − 乙太網路）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mkt']}-PL!{FC[y]}{A_['emk']}", "#,##0.0", "幾乎全為 NVIDIA"),
        ("sh", "Arista 有效份額（含前端與 scale-across）", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以公司 AI 目標校準"),
        ("ai", "Arista AI 營收（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}", "#,##0.00", ""),
        ("oth", "其他廠商（NVIDIA Spectrum-X、白牌、Cisco 等）＝ 乙太網路市場 − Arista", "US$B", lambda y: f"=PL!{FC[y]}{A_['emk']}-PL!{FC[y]}{A_['ai']}", "#,##0.0", "上限：Arista AI 含前端，所以此列偏低"),
        ("itr", "Arista AI 營收 ÷ AI IT 資本支出", "%", lambda y: f"=PL!{FC[y]}{A_['ai']}/itc_{y % 100}", "0.00%", "markets/network 的 AI IT 資本支出（加速器 × 每顆 IT capex）")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：FY2026 有效份額約 15%，與 Dell'Oro「Celestica＋NVIDIA 合計約 50%、Arista 第三」相容；本模型的份額路徑（相對值）是 Assumed，是最需要每季檢查的輸入。").font = BOLD
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」與十年 P/E 規則值。Arista 的 GAAP 與 non-GAAP 差距小（約 7%），十年 GAAP P/E 適用；規則值以 GAAP 換算 EPS（non-GAAP × g2n）計算"
V["A2"] = "循環位置：Arista 不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 NVIDIA；FY(t)＝CY(t)，不平移）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-23 收盤）"); V["B6"] = "=px"; V["B6"].number_format = "0.00"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "前瞻 P/E（StockAnalysis）"); V["B11"] = "=pe_fwd_sa"; V["B11"].number_format = "0.0"
V.cell(12, 1, "  隱含未來 12 月 EPS ＝ 現價 ÷ 前瞻 P/E"); V["B12"] = "=px/pe_fwd_sa"; V["B12"].number_format = "0.00"
V.cell(13, 1, "近 12 月 GAAP P/E（StockAnalysis）"); V["B13"] = "=pe_ttm_gaap"; V["B13"].number_format = "0.0"
V["E8"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"; V["E11"] = "Andy 先驗 A9a 答 30x（口徑未註明）"
V.cell(15, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, 16, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（GAAP P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 17):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(21, 1, "3. 規則值隱含股價（美元；約 FY2027 年底，未折現）＝ FY2027 GAAP 換算 EPS × 規則倍數").font = BOLD
hdr(V, 22, ["情境", "FY2027 non-GAAP EPS", "GAAP 換算 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ GAAP 換算 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 23 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}*g2n").number_format = "0.00"
    V.cell(i, 4, f"=D{17 + j}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(26, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（情境 A 的 FY2028 EPS，non-GAAP；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, 27, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，C）", "對照：Andy 合理 P/E", "對照：十年中位數（GAAP）"])
vr = 28
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, "=pe_fair_andy").number_format = "0.0"; V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。")
vr += 2
V.cell(vr, 1, "5. 對照：以 Andy 合理 P/E 25x 計算的隱含價（non-GAAP EPS；只作對照，非規則值）").font = BOLD; vr += 1
hdr(V, vr, ["情境", "FY2027 EPS", "× 25x", "相對現價", "FY2028 EPS", "× 25x", "相對現價"]); vr += 1
IMP = {}
for s in SC:
    V.cell(vr, 1, s)
    V.cell(vr, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 3, f"=B{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 4, f"=C{vr}/px-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 6, f"=E{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
    IMP[s] = vr; vr += 1
assert vr < 50, vr
PE_CMC = [36.2, 40.3, 24.8, 18.1, 34.8, 52.4, 27.5, 34.9, 48.6, 46.8]          # companiesmarketcap 年底 P/E 2016–2025
PE_MT = [38.86, 46.56, 26.91, 19.12, 36.37, 54.50, 28.42, 35.85, None, None]    # Macrotrends 12-31 P/E 2016–2023（頁面只到 2024-04）
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2016–2025 共 10 年）"},
    {"label": "Macrotrends（12 月 31 日 ÷ TTM GAAP EPS）", "vals": PE_MT, "tag": "Interested-party（資料商；2016–2023 共 8 年，2024 起頁面無資料：未找到）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 GAAP 換算 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 GAAP 換算 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 GAAP 換算 EPS（A）", "basis": f"=C{VAL['A']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位；兩個來源各取分位再平均。歷史倍數為 GAAP 落後 P/E，所以基準用 GAAP 換算 EPS。2016–2025 含高成長期（2021 年 52x），規則值可能高估成長放緩後的倍數。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額、非 AI 毛利率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("份額流失（低）", "份額相對值 2027 0.80、2028 0.70、2029 0.65（NVIDIA、白牌加速）", dict(rel="RLO")),
         ("份額持平（高）", "份額相對值 1.05", dict(rel="RHI")),
         ("乙太網路占比 低", "2026 70% → 2029 75%", dict(eth="ETLO")), ("乙太網路占比 高", "2026 80% → 2029 92%", dict(eth="ETHI")),
         ("AI 市場 ×0.8", "FY2027 起 scale-out 交換器市場 ×0.8（供給受限或內容不如預期）", dict(mf="mf_lo")), ("AI 市場 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("非 AI 成長 低", "2027 +3%、2028–29 +2%", dict(gn="GNLO")), ("非 AI 成長 高", "2027 +15%、2028 +12%、2029 +10%", dict(gn="GNHI")),
         ("AI 毛利率差距 3pt", "AI 毛利率只低 3 個百分點", dict(d="d_ai_lo")), ("AI 毛利率差距 10pt", "低 10 個百分點", dict(d="d_ai_hi")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi"))]
# 低／高路徑的具名範圍（逐年）：以 _lo／_hi 名稱組合
for nm, base, suf in (("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("ETLO", "eth", "_lo"), ("ETHI", "eth", "_hi"), ("GNLO", "g_nai", "_lo"), ("GNHI", "g_nai", "_hi")):
    for y in (26, 27, 28, 29):
        src = f"{base}_{y}{suf}" if base != "g_nai" else f"{base}{y}{suf}"
        if base == "rel" and y == 26: continue
        if base == "g_nai" and y == 26: continue
        target = f"{nm}_{y}" if base != "g_nai" else f"{nm}{y}"
        wb.defined_names[target] = DefinedName(target, attr_text=wb.defined_names[src].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.00"), ("D", "ai", 2027, "#,##0.00"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.00"), ("I", "ai", 2028, "#,##0.00"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數情境 B、C 的差異從 FY2028 起。份額與乙太網路占比的案例中，FY2026 有效份額依新的乙太網路占比重算，所以 FY2026 恆等於錨點。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、共識、研調份額。只作事後對照，不為對上數字而調參數。2027 年共識（StockAnalysis 付費）未找到，只有單一分析師與前瞻 P/E 隱含值"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與分析師（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "StockAnalysis", "錨點＝公司全年指引 12.6")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "StockAnalysis", "錨點（3Q、4Q 依指引與 In 頁假設）")
rec("FY2027 EPS vs Erste Group（單一分析師）", "=erste_eps27", pl("eps", 2027), "MarketBeat", "2027 共識：未找到（付費）")
rec("FY2027 EPS vs 前瞻 P/E 隱含未來 12 月 EPS", "=px/pe_fwd_sa", pl("eps", 2027), "StockAnalysis（推導）", "未來 12 月 ≈ 4Q26＋¾ FY2027，口徑不同，只作量級比較")
rec("FY2027 營收年增（模型）", None, pl("g", 2027), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "B. 份額與市場（先驗比對的機制檢查）").font = BOLD; rr_ += 1
rec("FY2026 Arista 有效份額 vs IDC 資料中心乙太網路份額（全部速率）", "=idc_anet", lambda s: f"=PL!D{PB[s]['sh0']}", "IDC 2Q26", "口徑不同：IDC 含前端與非 AI；本模型分母只含 AI 後端", fmt="0.0%", pt=True)
rec("FY2026 Arista 有效份額 vs 1 − Celestica、NVIDIA 合計（上限）", "=1-dell_top2", lambda s: f"=PL!D{PB[s]['sh0']}", "Dell'Oro 2025", "Arista 第三，有效份額應低於上限", fmt="0.0%", pt=True)
rec("FY2027 AI ÷ 營收 vs Andy 先驗（54%）", 0.54, pl("aish", 2027), "先驗_Andy_Arista A2（Drive）", "", fmt="0.0%", pt=True)
rec("FY2027 Arista AI 營收 ÷ AI IT 資本支出", None, lambda s: f"=PL!E{PB[s]['ai']}/itc_27", "markets/network", "Andy 先驗 A7：每 GW 4 億美元", fmt="0.00%")
rec("FY2028 Arista AI 營收 ÷ AI 後端乙太網路市場（有效份額）", None, pl("sh", 2028), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("非 AI 毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "AI 毛利率＝此值 − d_ai", fmt="0.0%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "AI 占比上升使毛利率下滑", fmt="0.0%", pt=True)
rec("3Q26 EPS（模型）vs 公司指引中點 1.07", 1.07, lambda s: "=Q!E15", "公司（2Q26 新聞稿）", "", fmt="0.000")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 Arista AI 營收 ÷ 部署 GW 計算每 GW 營收，對照 Andy 先驗 A7"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝Arista 8-K 新聞稿，已逐份抓取；法說內容只經二手轉述）"; SO["A1"].font = BOLD
srcs = [("公司：2Q26 新聞稿（含上半年合計、3Q26 指引）", PR_Q226),
        ("公司：4Q25 與 FY2025 新聞稿", PR_Q425),
        ("二手：BigGo Finance 2Q26 法說摘要（全年指引 12.6、AI 至少 3.5、scale-across 約 1.2、毛利率與營業利益率指引）", "https://finance.biggo.com/news/US_ANET_2026-08-04"),
        ("二手：sec-api 彙整（客戶集中、遞延收入、GAAP／non-GAAP 毛利率與營業利益率）", "https://sec-api.io/insights/arista-networks-anet-stock-analysis-q2-2026-ai-networking-revenue-deferred-revenue-customer-concentration-insider-sales-and-scenarios-to-fy2029"),
        ("二手：Next Platform（2025 營收、2025 AI 約 1.63、微軟 26%、Meta 16%）", "https://www.nextplatform.com/2026/02/13/the-current-ai-networking-wave-will-be-a-tsunami-of-money-by-2027/"),
        ("研調：IDC 2Q26 乙太網路交換器（NVIDIA 20.4%、Arista 18.7%、Cisco 18.2%）", "https://www.idc.com/resource-center/blog/ethernet-switch-market-surges-43-4-to-18-9b-in-2q26-as-ai-infrastructure-demand-drives-record-datacenter-spending/"),
        ("研調：Dell'Oro 2025 AI 後端（乙太網路超過 InfiniBand 兩倍；Celestica、NVIDIA 並列第一，Arista 第三）", "https://www.delloro.com/news/ethernet-more-than-doubles-size-of-infiniband-as-the-leading-fabric-for-ai-scale-out-networks-in-2025/"),
        ("二手：StockAnalysis 現價、前瞻 P/E、2026 共識", "https://stockanalysis.com/stocks/anet/forecast/"),
        ("二手：MarketBeat（Erste FY2027 EPS 4.79）", "https://www.marketbeat.com/instant-alerts/estimates-fy2027-eps-forecast-for-arista-networks-increased-by-analyst-2026-10-23/"),
        ("資料商：companiesmarketcap 年底 P/E", "https://companiesmarketcap.com/arista-networks/pe-ratio/"),
        ("資料商：Macrotrends 季度 P/E（至 2024-04）", "https://www.macrotrends.net/stocks/charts/ANET/arista-networks/pe-ratio"),
        ("市場層：markets/network、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：2027 共識（付費）；各季 AI 營收；產品別毛利率；前瞻 P/E 十年歷史（Bloomberg）；重編後 FY2025 non-GAAP 淨利", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Arista Networks（ANET）收支與估值模型 v0.1（AI 半導體 Project：網通系統端樣板）　2026-10-24", True),
 ("版本：v0.1（初版）。Arista 在 AI 後端乙太網路交換器的份額在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）。", False),
 ("模型命題：Arista 的 AI 營收＝AI 後端 scale-out 交換器市場（本專案網通模型）× 乙太網路占比 × Arista 份額（含前端與 scale-across 附掛）；非 AI（雲端前端、企業、園區）另依成長率推。股價約 37 倍前瞻 EPS：在 NVIDIA（Spectrum-X）與白牌夾擊下，Arista 的份額路徑能否支撐這個倍數？", False),
 ("核心問題：2027–28 年 Arista 在 AI 乙太網路的份額每掉一截，EPS 少多少；現價隱含的份額路徑合不合理。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029（三年）。口徑：公司 non-GAAP；十年 P/E 規則值用 GAAP 換算 EPS。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 營收：markets/network 的 AI 後端 scale-out 交換器市場 × 情境顆數係數 × 乙太網路占比（Assumed）× Arista 有效份額（FY2026 以公司 AI 目標 3.5 校準）× 份額相對值（Assumed）。", False),
 ("非 AI 營收：FY2026 錨點（營收指引 − AI 目標）× Assumed 成長率。毛利率：AI 比非 AI 低 d_ai（Assumed），非 AI 毛利率為 FY2026 殘差（Derived）。營業費用：彈性 phi（Derived）。", False),
 ("校準：FY2026 ＝ 1Q–2Q 實績＋3Q 指引＋全年指引（營收 12.6、毛利率 63%、營業利益率 48.5%）。不用共識（只在 Recon 對帳）。", False),
 ("估值：現價 ÷ FY2027、FY2028 EPS；十年 GAAP P/E 規則值（兩個資料商來源）× FY2027 GAAP 換算 EPS；循環位置以本模型營收最高年判定（Arista 不套設備股見頂年）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文＝SEC EDGAR 8-K；法說內容只經二手轉述（標二手）。", False),
 ("", False),
 ("已知限制：(1) 份額相對值與乙太網路占比皆為 Assumed，是 FY2028 EPS 最敏感的輸入（Sens）；(2) FY2026 AI 用公司目標「至少 3.5」，實際可能較高；(3) 網通模型 2027 年 scale-out 交換器 +90%，公司稱供給吃緊到 2028 年，模型未設供給上限（Sens 用市場 ×0.8）；(4) 2027 共識未找到，只有單一分析師；(5) 十年 P/E 含高成長期溢價；(6) R4 產業部署路徑未建（Recon 預留列）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("ARISTA_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK_ROW, "SENSN": SENS_N, "IMP": IMP, "VAL": VAL, "SB": SB},
              open(os.environ["ARISTA_ROWMAP"], "w"), ensure_ascii=False, indent=1)
