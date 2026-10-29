# 美超微 Super Micro Computer（SMCI US）收支與估值模型 v0.1.1（AI 半導體 Project：伺服器組裝次產業，樣板＝wiwynn/、quanta/）
# 用法：python3 smci/build_smci.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度 FY(t)＝t−1 年 7 月～t 年 6 月（6 月底結束）。上游市場層為日曆年：FY(t) 市場＝½ × CY(t−1)＋½ × CY(t)（半年加權）。
# 金額：十億美元（US$B）；EPS：美元（GAAP 稀釋）。口徑：歸屬母公司淨利＝扣除特別股股利後、歸屬普通股的 GAAP 淨利。
# FY2026（2025-07～2026-06）已結束：FY2026 為實績錨點；FY2027–FY2029 由驅動因子推導；FY2030–FY2032 為收斂法。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith("um") or str(k).startswith("u_") or str(k).startswith("g29_")):
        TS[k] = (r[4].value, r[0].row)
PLATS = ("bw", "rb", "ru", "hp", "amd", "tpu", "trn", "oth")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
assert all(f"u_{p}_{y}" in TS for p in PLATS for y in (26, 27, 28)) and all(f"g29_{p}" in TS for p in PLATS), TS.keys()
def units(p, y):
    return TS[f"u_{p}_{y}"][0] if y < 29 else TS[f"u_{p}_28"][0] * (1 + TS[f"g29_{p}"][0])
# 商用 GPU（美超微可銷售的平台）：NVIDIA Blackwell、Rubin、Rubin Ultra、Hopper 與 AMD Instinct；不含 Google TPU、Trainium、其他 ASIC
# 世代分組：GB300 世代（Blackwell、Hopper；AMD 2026 的 MI355）、Rubin 世代（Rubin、Rubin Ultra；AMD 2027 起的 MI450 整櫃，Analogy）
OLD = {y: sum(units(p, y) for p in ("bw", "hp")) + (units("amd", y) if y == 26 else 0) for y in (26, 27, 28, 29)}
NEW = {y: sum(units(p, y) for p in ("rb", "ru")) + (units("amd", y) if y >= 27 else 0) for y in (26, 27, 28, 29)}
ALL = {y: sum(units(p, y) for p in PLATS) for y in (26, 27, 28, 29)}
U_SRC = f"{rel(TSMC_F)}｜In!E{TS['u_bw_26'][1]}:E{TS['g29_oth'][1]}"
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("rack", "gpu_n"):
        BR[k] = (r[4].value, r[7].value, f"{rel(BR_F)}｜In!E{r[0].row}（GB300）、H{r[0].row}（VR200）")
assert len(BR) == 2, BR.keys()
# CF10：商用 GPU 系統市場只有一個值（Siliconomics Rack 頁）；本模型的世代分組與每顆系統價值必須重現同一市場（情境 A）
for j, y in enumerate((26, 27, 28, 29)):
    _m = OLD[y] * BR["rack"][0] / BR["gpu_n"][0] * 1000 + NEW[y] * BR["rack"][1] / BR["gpu_n"][1] * 1000
    assert abs(_m - SIL.NAMES["SIF_GPUSysMkt_A"][j]) < 1e-6, ("商用 GPU 系統市場與 Siliconomics 不一致", y, _m, SIL.NAMES["SIF_GPUSysMkt_A"][j])

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；每顆系統價值：千美元；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個。EPS 為 GAAP 稀釋、美元"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司原文＝美超微新聞稿（8-K）、財報說明稿、10-K（經二手摘要）；公司的 AI 占比、指引與展望為 Interested-party；券商與資料商估計皆為 Interested-party。商用 GPU 顆數與情境係數讀台積電模型、每顆系統價值讀橋接表（唯讀）"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
PR4Q = "https://ir.supermicro.com/news/news-details/2026/Supermicro-Announces-Fourth-Quarter-and-Full-Fiscal-Year-2026-Financial-Results/default.aspx"
PREP = "https://s204.q4cdn.com/707617056/files/doc_financials/2026/q4/V2/FQ426-SMCI-Prepared-Remarks.pdf"
MB0811 = "https://www.marketbeat.com/instant-alerts/super-micro-computer-q4-earnings-call-highlights-2026-08-11/"
PRQ2 = "https://www.sec.gov/Archives/edgar/data/1375365/000137536526000006/exhibit991_20251231.htm"
PRQ1U = "https://ir.supermicro.com/news/news-details/2026/Supermicro-Provides-First-Quarter-of-Fiscal-Year-2027-Preliminary-Business-Update/default.aspx"
INVEST = "https://www.supermicro.com/en/pressreleases/supermicro-announces-completion-independent-investigation-and-continued-enhancement"
ST0820 = "https://stocktwits.com/news-articles/markets/equity/smci-stock-gains-supermicro-clears-senior-management-export-control-case/cZYIq3bRJZ5"
FIN0611 = "https://ir.supermicro.com/news/news-details/2026/Supermicro-Announces-Pricing-of-Equity-and-Equity-Linked-Financing-Transactions-To-Fund-AI-Orders/default.aspx"
K10 = "https://www.sec.gov/Archives/edgar/data/0001375365/000137536526000022/smci-20260630.htm"
STK10 = "https://stockti.com/super-micro-revenue-surges-78-but-margin-pressures-and-control-gaps-linger"
FUT = "https://futurumgroup.com/insights/supermicro-q2-fy2026-delivers-breakout-ai-gpu-platform-revenue"
SA = "https://stockanalysis.com/stocks/smci/statistics/"
SAF = "https://stockanalysis.com/stocks/smci/forecast/"
SAR = "https://stockanalysis.com/stocks/smci/financials/ratios/"
SAI = "https://stockanalysis.com/stocks/smci/financials/"
YF = "https://finance.yahoo.com/quote/SMCI/analysis/"
TICK = "https://tickflow.io/stock/SMCI/forecast"
SWS = "https://simplywall.st/stocks/us/tech/nasdaq-smci/super-micro-computer/future"
MT = "https://www.macrotrends.net/stocks/charts/SMCI/super-micro-computer/pe-ratio"
MTNI = "https://www.macrotrends.net/stocks/charts/SMCI/super-micro-computer/net-income"
MTREV = "https://www.macrotrends.net/stocks/charts/SMCI/super-micro-computer/revenue"
FR = "https://fullratio.com/stocks/nasdaq-smci/pe-ratio"
P_DELL = "https://stockanalysis.com/stocks/dell/statistics/"
P_HPE = "https://stockanalysis.com/stocks/hpe/statistics/"
P_WY = "https://stockanalysis.com/quote/tpe/6669/statistics/"
P_QC = "https://stockanalysis.com/quote/tpe/2382/statistics/"
P_WS = "https://stockanalysis.com/quote/tpe/3231/statistics/"
URLNAME = {PR4Q: "公司 FY2026 第四季新聞稿（2026-08-11）", PREP: "公司 4Q FY26 財報說明稿", MB0811: "MarketBeat 2026-08-11（法說重點）",
           PRQ2: "公司 2Q FY26 新聞稿（8-K）", PRQ1U: "公司 1Q FY27 初步營運更新（2026-10-20）", INVEST: "公司 2026-08 獨立調查完成公告",
           ST0820: "Stocktwits 2026-08-20", FIN0611: "公司 2026-06-11 股權與特別股定價公告", K10: "公司 FY2026 10-K（2026-08-31）", STK10: "StockTi（10-K 摘要）",
           FUT: "Futurum（2Q FY26）", SA: "stockanalysis 統計頁", SAF: "stockanalysis 預測頁", SAR: "stockanalysis 比率頁", SAI: "stockanalysis 損益表",
           YF: "Yahoo Finance 分析頁", TICK: "tickflow 預測頁", SWS: "Simply Wall St", MT: "Macrotrends P/E", MTNI: "Macrotrends 淨利", MTREV: "Macrotrends 營收",
           FR: "fullratio P/E", P_DELL: "stockanalysis（Dell）", P_HPE: "stockanalysis（HPE）", P_WY: "stockanalysis（緯穎）", P_QC: "stockanalysis（廣達）", P_WS: "stockanalysis（緯創）"}

sec("A. 商用 GPU 顆數與每顆系統價值（讀 tsmc、markets/chip_bridge 成品；情境 A 基準，B／C 依顆數係數縮放）")
for y in (26, 27, 28, 29):
    inp(f"uo_{y}", f"GB300 世代商用 GPU 顆數 CY20{y}（Blackwell＋Hopper" + ("＋AMD MI355）" if y == 26 else "）"), "百萬顆", None, round(OLD[y], 6), None, "Derived",
        U_SRC if y == 26 else "同上", "台積電模型八類加速器中 NVIDIA Blackwell、Hopper（2026 另含 AMD）；2029＝2028 × (1＋各類成長)" if y == 26 else "同上")
for y in (26, 27, 28, 29):
    inp(f"un_{y}", f"Rubin 世代商用 GPU 顆數 CY20{y}（Rubin＋Rubin Ultra" + ("）" if y == 26 else "＋AMD MI450）"), "百萬顆", None, round(NEW[y], 6), None, "Derived",
        U_SRC if y == 26 else "同上", "AMD 自 2027 年起改計 Rubin 世代價值（MI450 整櫃；Analogy）" if y == 27 else "同上")
for y in (26, 27, 28, 29):
    inp(f"ut_{y}", f"全部 AI 加速器顆數 CY20{y}（對照）", "百萬顆", None, round(ALL[y], 6), None, "Derived", U_SRC if y == 26 else "同上", "八類合計；商用 GPU 以外為 TPU、Trainium、其他 ASIC（美超微不銷售）" if y == 26 else "同上")
inp("v_gb", "每顆系統價值：GB300 世代（整櫃買方價 ÷ 每架 GPU 封裝數）", "千美元/顆", None, f"={BR['rack'][0]}/{BR['gpu_n'][0]}*1000", None, "Derived", BR["rack"][2],
    f"Tokenomics v20.1 機架單價 {BR['rack'][0]} 百萬美元 ÷ {BR['gpu_n'][0]} 顆（經橋接表；唯讀）；HGX 八卡伺服器與 AMD MI355 以同值近似（Analogy）")
inp("v_vr", "每顆系統價值：Rubin 世代（VR200 整櫃買方價 ÷ 每架 GPU 封裝數）", "千美元/顆", None, f"={BR['rack'][1]}/{BR['gpu_n'][1]}*1000", None, "Derived", BR["rack"][2],
    f"機架單價 {BR['rack'][1]} 百萬美元 ÷ {BR['gpu_n'][1]} 顆；Rubin Ultra 與 AMD MI450 未另訂（Analogy：同 VR200）")
inp("vpg_28", "每顆系統價值 2028 ÷ 2027（Rubin 世代）", "x", 0.85, 1.0, 1.2, "Assumed", "Rubin Ultra（Kyber）整櫃價格：未找到", "與廣達、緯穎模型的每 GW 價值變動同區間")
inp("vpg_29", "每顆系統價值 2029 ÷ 2027（Rubin 世代）", "x", 0.85, 1.0, 1.2, "Assumed", "同上", "—")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", f"{rel(TSMC_F)}｜In!E{TS[f'um28_{s}'][1]}", "A＝順利（基準）；B＝平平（AI 資本支出消化）；C＝不如預期（明顯下修）；與台積電、廣達、緯穎模型相同")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", f"{rel(TSMC_F)}｜In!E{TS[f'um29_{s}'][1]}", "同上")
inp("mf", "AI 部署係數（Sens 用；CY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "台積電模型 CY2027 商用 GPU 顆數約 1.4 倍（8.35 → 11.2 百萬顆）", "±20%")

sec("B. 營收結構、份額（FY2027 以公司指引校準）")
inp("ais_25", "FY2025 AI GPU 平台 ÷ 營收", "%", 0.65, 0.70, 0.75, "Interested-party", "公司：FY2025 AI GPU 平台占營收逾七成（法說；原文未逐字取得）", "量化第二來源：未找到")
inp("ais_1h", "1H FY26（2025-07～12）AI GPU 平台 ÷ 營收", "%", 0.75, 0.85, 0.90, "Assumed", FUT + "：2Q FY26 營收由 GB300 NVL72 與 AMD MI355 量產帶動（未給比例）", "公司揭露比例：未找到")
inp("ais_q3", "3Q FY26（2026-01～03）AI ÷ 營收", "%", 0.78, 0.80, 0.85, "Interested-party", PREP + "：3Q 逾 80%（公司）", MB0811 + "：同")
inp("ais_q4", "4Q FY26（2026-04～06）AI ÷ 營收", "%", 0.55, 0.60, 0.65, "Interested-party", PREP + "：約 60%，專案出貨時點使比重下降（公司）", MB0811 + "：同；公司預期未來逾 80% 營收與 AI 相關、純 AI 約 60–70%")
inp("g27", "FY2027 營收指引（情境 A＝上緣、B＝中點、C＝下緣）", "US$B", 65.0, 68.5, 72.0, "Interested-party", PR4Q + "：FY2027 營收 650～720 億美元（公司，2026-08-11）", PRQ1U + "：1Q FY27 初步營收 185～195 億美元，高於原指引 145～155 億；全年指引未調整（公司，2026-10-20）")
inp("rel_28", "美超微份額 FY2028 ÷ FY2027（A＝高、B＝基準、C＝低）", "x", 0.85, 0.95, 1.0, "Assumed", "Dell、HPE 與 ODM 直供（廣達、緯創、鴻海）爭取 neocloud 與企業訂單；4Q FY26 逾 600 億美元新訂單（公司）", "最大客戶占 FY2026 營收 28%（公司說明稿）：大客戶轉單為下檔")
inp("rel_29", "美超微份額 FY2029 ÷ FY2027", "x", 0.80, 0.92, 1.0, "Assumed", "同上", "—")
inp("g_nai27", "非 AI（企業、儲存、邊緣）營收年增 FY2027", "%", 0.0, 0.15, 0.30, "Assumed", PREP + "：企業與通路營收 4Q FY26 年增 172%（含 AI，公司）", "非 AI 營收分項：未找到；FY2026 非 AI 為殘差（Q 頁）")
inp("g_nai28", "非 AI 營收年增 FY2028", "%", -0.05, 0.08, 0.20, "Assumed", "同上", "—")
inp("g_nai29", "非 AI 營收年增 FY2029", "%", -0.05, 0.06, 0.15, "Assumed", "同上", "—")

sec("C. 毛利率（A 最有利、C 最不利）")
inp("gm_nai", "非 AI 營收毛利率", "%", 0.13, 0.15, 0.18, "Assumed", "FY2023（AI 占比低）GAAP 毛利率約 18%（公司年報，stockanalysis 損益表：毛利 1,283 ÷ 營收 7,123）", "部門別毛利率：未找到；公司 DCBBS 目標毛利率 20%＋（Futurum 轉述）")
inp("dgm_27", "FY2027 AI 毛利率 − FY2026 校準值（A＝高、B＝基準、C＝低）", "pt", -0.005, 0.005, 0.015, "Interested-party", PRQ1U + "：1Q FY27 初步毛利率 12.0%–12.5%，客戶與產品組合（公司）", PREP + "：4Q FY26 17.6%（non-GAAP）約四分之三來自組合與遞延至 FY27 的合約、四分之一來自關稅與存貨跌價回轉（可能不重複）")
inp("dgm_28", "FY2028–FY2029 AI 毛利率 − FY2026 校準值", "pt", -0.015, 0.003, 0.015, "Assumed", "C：Dell、ODM 價格競爭與 GPU 轉手比重上升；A：液冷與 DCBBS 整體解決方案", "FY2024 13.8% → FY2025 11.1% → FY2026 10.8%（GAAP；stockanalysis 損益表）")

sec("D. 1Q FY27 初步數字（Recon 對帳用）")
inp("q1_rev", "1Q FY27（2026-07～09）初步營收", "US$B", 18.5, 19.0, 19.5, "Interested-party", PRQ1U + "：185～195 億美元（公司，未經審計）", SWS + "：同（2026-10）")
inp("q1_gm", "1Q FY27 初步毛利率", "%", 0.120, 0.1225, 0.125, "Interested-party", PRQ1U + "：12.0%–12.5%（公司）", "原指引 10.4%–10.8%（4Q FY26 財報說明稿）")

sec("E. 營業費用、淨利係數、特別股與股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.3, 0.4, 0.6, "Derived", PREP + "：1Q FY27 GAAP 營業費用約 4.53 億美元（年化約 18.1 億，較 FY2026 14.57 億 +24%）", "FY2026 實測：營業費用 +24%、營收 +78%（Q 頁），彈性約 0.3")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.002, 0.0, 0.002, "Assumed", "無", "營益率約 8%，±0.2 個百分點約 ±2.5% 的營業利益")
inp("nio_f", "淨利（特別股股利前）÷ 營業利益", "%", 0.68, 0.75, 0.80, "Derived", "FY2026 2,230.5 ÷ 2,770.5＝0.805；稅率 FY2026 19.9%、1Q FY27 指引 20.1%（公司）", "1Q FY27 指引隱含約 0.68（GAAP EPS 0.89–0.98 × 7.45 億股＋特別股股利，÷ 營收 150 億 × 毛利率 10.6% − 營業費用 4.53 億）：銀行借款 40.6 億、可轉債 46.6 億美元的利息（Derived）")
inp("pref_div", "強制可轉換特別股股利（年）", "US$B", None, "=86.25*50/1000*0.07", None, "Derived", FIN0611 + "：7,500 萬股存託股＋超額配售 1,125 萬股、每股 50 美元、股利率 7.0%，2029-06-01 強制轉換（公司）", PR4Q + "：4Q FY26 募得淨額 56 億美元，其中強制可轉換特別股 42 億（公司）；超額配售全數執行為 Assumed")
inp("sh27", "FY2027 GAAP 稀釋股數", "十億股", None, 0.745, None, "Interested-party", PR4Q + "：1Q FY27 指引 GAAP 稀釋股數 7.45 億（公司）", SA + "：流通股數 6.5697 億（2026-10-24）")
inp("sh_g", "股數年變動（員工股票報酬、可轉債）", "%", 0.0, 0.01, 0.03, "Assumed", "FY2025 6.284 → FY2026 6.973 億股（GAAP 稀釋，含 6 月增資；公司）", "1Q FY27 指引股票報酬費用約 1.06 億美元／季（公司）")
inp("conv_sh", "特別股轉換新增股數（FY2030 起；最低～最高轉換率中點）", "十億股", "=86.25*1.5152/1000", "=86.25*(1.5152+1.8182)/2/1000", "=86.25*1.8182/1000", "Derived", FIN0611 + "：每股存託股轉換 1.5152～1.8182 股普通股（公司）", "轉換後特別股股利消失；FY2029 股利以全年計（實際 2029-06 轉換，差約一個月）")
inp("sbc_ps", "GAAP → non-GAAP EPS 調整（每股；Recon 對共識用）", "US$", None, "=(0.106-0.032)*4/sh27", None, "Derived", PR4Q + "：1Q FY27 股票報酬約 1.06 億美元、稅效果 0.32 億（公司）", "FY2026 non-GAAP 3.63 − GAAP 3.26＝0.37（公司）")

sec("F. 共識（只作 Recon 對帳，不作校準；共識 EPS 為 non-GAAP）")
inp("c_rev27", "共識 FY2027 營收", "US$B", 63.5, 70.0, 82.0, "Interested-party", SAF + "：平均 700 億美元（15 位，低 635、高 820；2026-10-22）", YF + "：同 700 億（15 位）")
inp("c_rev28", "共識 FY2028 營收", "US$B", 70.84, 81.4, 92.0, "Interested-party", YF + "：平均 814 億美元（15 位）", "stockanalysis FY2028 需付費：未找到")
inp("c_eps27", "共識 FY2027 EPS（non-GAAP）", "US$", 3.50, 4.67, 6.40, "Interested-party", SAF + "：4.67（15 位，低 3.50、高 6.40）", YF + "：4.67（18 位）")
inp("c_eps28", "共識 FY2028 EPS（non-GAAP）", "US$", None, 5.59, None, "Interested-party", YF + "：5.59（17 位）", "第二來源：未找到（tickflow 季度數字前後不一致，不採用）")

sec("G. 價格與市場倍數")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 32.94, None, "Interested-party", SA + "：10-23 收盤 32.94", FR + "：10-23 價格 32.94、近 12 月 EPS 3.65")
inp("pe_ttm", "近 12 月 P/E（stockanalysis，10-23）", "x", None, 10.12, None, "Interested-party", SA, "fullratio 9.02（以 non-GAAP EPS）")
inp("fpe", "前瞻 P/E（stockanalysis，10-23）", "x", None, 7.06, None, "Interested-party", SA, "以 non-GAAP 共識 EPS")
inp("cons_feps", "共識前瞻 12 月 EPS（10-23 價格 ÷ 前瞻 P/E；non-GAAP）", "US$", None, "=px/fpe", None, "Derived", "stockanalysis 前瞻 P/E 7.06（2026-10-23）", "≈ 共識 FY2027 EPS 4.67")
inp("cons_tp", "市場共識目標價（賣方平均）", "US$", 15.0, 43.27, 60.0, "Interested-party", SAF + "：18 位平均 43.27、中位數 45、最低 15、最高 60（S&P Global，2026-10-22）", TICK + "：25 位平均 41.44（2026-10-24）；Simply Wall St 08-12：41.75（16 位）")
REV10 = [2225, 2485, 3360, 3500, 3339, 3557, 5196, 7123, 14989, 21972]
NI10 = [72, 67, 46, 72, 84, 112, 285, 640, 1153, 1049]
inp("nim10", "十年平均淨利率（FY2016–FY2025，淨利 ÷ 營收）", "%", None, "=AVERAGE(" + ",".join(f"{n}/{v}" for n, v in zip(NI10, REV10)) + ")", None, "Derived",
    MTREV + "、" + MTNI + "：各年營收與淨利（百萬美元；FY2016 2,225／72 … FY2023 7,123／640）", SAI + "：FY2024 14,989／1,153、FY2025 21,972／1,049")

sec("H. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%；美、台、日、韓同一值）", "DCF 折現率；DCF 情境價值 ×(1＋r) 推到 12 個月後。治理與會計風險不另加折現率（見 Sens、README）")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "DCF 永續成長率；FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均，FY2024–26）", "%", None, "=AVERAGE(0.312,0.179,0.2147)", None, "Derived", SAR + "：ROE FY2024 31.20%、FY2025 17.90%、FY2026 21.47%", "單一來源；Simply Wall St 三年後預測 22.3%")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（Andy 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2029 特別股股利前淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）；FY2030 起特別股已轉換，股利消失", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（情境 B FY2029 與十年平均的中點）", "%", None, 0, None, "Derived", "PL 頁情境 B FY2029、In 頁 nim10", "—")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（＝十年平均淨利率）", "%", None, "=nim10", None, "Derived", "In 頁 nim10", "—")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    for u_, nm_ in URLNAME.items():          # 網址只列在 Sources 頁；In 頁寫來源名稱
        s1 = s1.replace(u_, nm_); s2 = s2.replace(u_, nm_)
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for k in TP_KEYS:
    for c in range(1, 10): IN.cell(INROW[k], c).fill = KEY
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億美元；EPS 為美元、GAAP 稀釋；股數為十億股）。藍字＝輸入；黑字＝公式。會計年度 FY2026＝2025-07～2026-06。實績取自美超微新聞稿（8-K）與財報說明稿"
Q["A2"] = "FY2026 錨點＝全年實績（公司 2026-08-11 公布，10-K 2026-08-31 申報；新聞稿標示未經審計，需待審計確認）。3Q 營收、營業費用、營業利益、淨利＝全年 − 1H − 4Q（Derived）；毛利＝營業利益＋營業費用。股數＝淨利 ÷ 公司揭露稀釋 EPS（可轉債加回使其與揭露稀釋股數略有差異）"
hdr(Q, 3, ["項目", "FY2025A", "1H FY26（2025-07～12）", "3Q FY26（2026-01～03）", "4Q FY26（2026-04～06）", "FY2026A（錨點）", "標記", "來源1（公司）", "來源2／說明"])
S1 = "公司 FY2026 第四季新聞稿（2026-08-11）；1H：公司 2Q FY26 新聞稿（8-K）"
S2 = "stockanalysis 損益表：FY2026 營收 39,063、毛利 4,227、營業利益 2,770（百萬美元）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "營收", {"B": 21.972, "C": 17.7003, "D": "=39.063-C4-E4", "E": 11.12}, tag="Verified／Derived（3Q）", s2="1Q FY26 5.0178、2Q 12.6825（8-K）；FY2026 39.063（stockanalysis）")
qrow(5, "毛利（＝營業利益＋營業費用）", {c: f"={c}8+{c}7" for c in "BCDE"}, tag="Derived", s2="GAAP 毛利率：FY2025 11.1%、1H 7.2%、3Q 9.9%、4Q 17.5%、FY2026 10.8%（公司）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(7, "營業費用（GAAP）", {"B": 1.1769, "C": 0.6094, "D": "=1.4568-C7-E7", "E": 0.4546}, tag="Verified／Derived（3Q）", s2="FY2026 1.4568（公司）")
qrow(8, "營業利益", {"B": 1.253, "C": 0.6566, "D": "=2.7705-C8-E8", "E": 1.4881}, tag="Verified／Derived（3Q）", s2="FY2026 2.7705（公司）")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="—")
qrow(10, "歸屬母公司淨利（GAAP）", {"B": 1.0489, "C": 0.5688, "D": "=2.2305-C10-E10", "E": 1.1782}, tag="Verified／Derived（3Q）", s2="3Q 4.83 億（公司說明稿）；stockanalysis FY2026 淨利 2,196（可能已扣特別股股利，未採用）")
qrow(11, "淨利 ÷ 營業利益（含業外、稅）", {c: f"={c}10/{c}8" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(14, "  公司揭露稀釋 EPS（GAAP）", {"B": 1.68, "C": 0.86, "D": 0.72, "E": 1.62}, fmt="0.00", ann="=3.26", tag="Verified", s2="non-GAAP：FY2025 2.06、FY2026 3.63、4Q 1.70（公司）")
qrow(12, "股數（十億股；＝淨利 ÷ 揭露稀釋 EPS）", {c: f"={c}10/{c}14" for c in "BCDE"}, fmt="0.0000", ann="=F10/F14", tag="Derived", s2="淨利 ÷ 揭露 EPS；FY2026 揭露稀釋股數 6.973 億（第 15 列）")
qrow(13, "EPS（GAAP 稀釋）", {c: f"={c}10/{c}12" for c in "BCDE"}, fmt="0.00", ann="=F10/F12", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(15, "  公司揭露稀釋股數（對照）", {"B": 0.6284, "C": 0.6741, "D": 0.692, "E": 0.705}, fmt="0.0000", ann="=0.6973", tag="Verified", s2="4Q 7.050 億、3Q 6.92 億（公司說明稿）")
qrow(16, "AI GPU 平台營收（＝營收 × AI 占比）", {"B": "=B4*ais_25", "C": "=C4*ais_1h", "D": "=D4*ais_q3", "E": "=E4*ais_q4"}, tag="Derived／Interested-party", s1="公司比例（In 頁 B 節）", s2="AI 營收金額未揭露：未找到")
qrow(17, "非 AI（企業、儲存、邊緣等）", {c: f"={c}4-{c}16" for c in "BCDE"}, tag="Derived", s1="—", s2="—")
Q.cell(19, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(20, 1, "FY2026 營收年增"); Q["B20"] = "=F4/B4-1"
Q.cell(21, 1, "FY2026 營業費用年增"); Q["B21"] = "=F7/B7-1"
Q.cell(22, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的對照）"); Q["B22"] = "=B21/B20"
Q.cell(23, 1, "FY2026 AI 占營收（Derived；1H、3Q、4Q 加權）"); Q["B23"] = "=F16/F4"
Q.cell(24, 1, "FY2026 毛利率（不含 4Q；1H＋3Q）"); Q["B24"] = "=(C5+D5)/(C4+D4)"
Q.cell(25, 1, "2026-01～06（2H FY26）AI GPU 平台營收（份額校準用）"); Q["B25"] = "=D16+E16"
Q.cell(26, 1, "FY2026 EPS 自算 − 公司揭露（應為 0）"); Q["B26"] = "=F13-F14"
for rr in range(20, 27): Q[f"B{rr}"].number_format = "0.0%" if rr not in (22, 25, 26) else "0.000"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 10), ("sh", 12), ("eps", 13), ("ai", 16), ("nai", 17)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$F${rr}")
defname("a26_aih2", "Q!$B$25")
Q.column_dimensions["A"].width = 60
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 70; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I", 2032: "J"}
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
CYS = (2026, 2027, 2028, 2029); F3 = (2027, 2028, 2029)
SCN_P = {"A": dict(g27="g27_hi", rel="rel_{y}_hi", d27="dgm_27_hi", d28="dgm_28_hi"),
         "B": dict(g27="g27", rel="rel_{y}", d27="dgm_27", d28="dgm_28"),
         "C": dict(g27="g27_lo", rel="rel_{y}_lo", d27="dgm_27_lo", d28="dgm_28_lo")}
DEFP = dict(scn="A", vpg="vpg_{y}", mf="mf", gn="g_nai{y}", gmn="gm_nai", opxa="0", nio="nio_f", phi="phi", fg="fade_g{s}", fm="fade_m{s}", rel_mult="1")
def engine(ws, r0, title, **over):
    s = over.get("scn", "A")
    P = dict(DEFP); P.update(SCN_P[s]); P.update(over)
    p = lambda k, y=None: P[k].format(y=(y % 100 if y else ""), s=s)
    um = lambda y: {2026: "", 2027: "", 2028: f"*um28_{s}", 2029: f"*um29_{s}"}[y]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026A（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 11, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    mfx = lambda y: "" if y == 2026 else f"*{P['mf']}"
    sec_("A. 商用 GPU 系統市場（D–G 欄的「CY」列為日曆年 CY2026–CY2029；「FY」列為美超微會計年度）")
    row("uo", "GB300 世代商用 GPU 顆數（CY）", "百萬顆", {c_(y): f"=uo_{y % 100}{um(y)}{mfx(y)}" for y in CYS}, fmt="0.00",
        note="Blackwell＋Hopper（2026 另含 AMD MI355）；台積電模型顆數 × 情境係數")
    row("un", "Rubin 世代商用 GPU 顆數（CY）", "百萬顆", {c_(y): f"=un_{y % 100}{um(y)}{mfx(y)}" for y in CYS}, fmt="0.00",
        note="Rubin＋Rubin Ultra（2027 起另含 AMD MI450）")
    row("vo", "每顆系統價值：GB300 世代（CY）", "千美元/顆", {c_(y): "=v_gb" for y in CYS}, fmt="0.0")
    row("vn", "每顆系統價值：Rubin 世代（CY）", "千美元/顆", {"D": "=v_vr", "E": "=v_vr", **{c_(y): f"=v_vr*{p('vpg', y)}" for y in (2028, 2029)}}, fmt="0.0",
        note="整櫃買方價 ÷ 每架 GPU 封裝數（橋接表）")
    row("mcy", "商用 GPU 系統市場（CY）", "US$B", {c_(y): f"=({c_(y)}{R['uo']}*{c_(y)}{R['vo']}+{c_(y)}{R['un']}*{c_(y)}{R['vn']})" for y in CYS}, fmt="#,##0.0",
        note="顆數 × 每顆系統價值；含 NVIDIA 與 AMD，不含 TPU、Trainium 與其他 ASIC")
    row("mfy", "商用 GPU 系統市場（FY＝½ 前一 CY＋½ 當年 CY）", "US$B", {c_(y): f"=0.5*{c_(y - 1)}{R['mcy']}+0.5*{c_(y)}{R['mcy']}" for y in F3}, fmt="#,##0.0",
        note="會計年度 7 月～6 月：各取前後日曆年一半（Derived；未考慮年內出貨季節性）")
    row("mfy0", "FY2027 市場（部署係數 1.0；份額校準用）", "US$B", {"E": f"=0.5*D{R['mcy']}+0.5*(uo_27*v_gb+un_27*v_vr)"}, fmt="#,##0.0",
        note="份額以基準市場校準，Sens 的部署係數才會改變 FY2027 營收")
    row("sh0", "美超微有效份額（2H FY26 校準）＝ 2026-01～06 AI 營收 ÷ ½ CY2026 市場", "%", {"D": f"=a26_aih2/(0.5*D{R['mcy']})"}, fmt="0.00%",
        note="Derived：CY2026 上下半年各半（Rubin 下半年放量，上半年市場偏小，此份額偏低估）；只作對照")
    row("g27", "FY2027 營收（公司指引；A 上緣、B 中點、C 下緣）", "US$B", {"E": f"={p('g27')}"}, fmt="#,##0.0", note="Interested-party：公司 2026-08-11 指引 650～720 億美元")
    row("rel", "份額相對值（÷ FY2027）", "x", {"E": "=1", **{c_(y): f"={p('rel', y)}*{P['rel_mult']}" for y in (2028, 2029)}}, fmt="0.00",
        note="A 最有利、C 最不利（Assumed）；治理事件案例另見 Sens")
    row("sh", "美超微有效份額（FY2027 以指引校準）", "%", {"E": f"=(E{R['g27']}-E{{nai}})/E{R['mfy0']}", **{c_(y): f"=$E{r}*{c_(y)}{R['rel']}" for y in (2028, 2029)}}, fmt="0.00%",
        note="FY2027：（指引營收 − 非 AI 營收）÷ 基準市場；4Q FY26 逾 600 億美元新訂單使份額高於 2H FY26")
    row("shr", "FY2027 份額 ÷ 2H FY26 份額", "x", {"E": f"=E{R['sh']}/D{R['sh0']}"}, fmt="0.00", note="Derived：指引隱含的份額上升幅度")
    row("ai", "AI GPU 平台營收", "US$B", {"C": "=a25_ai", "D": "=a26_ai", **{c_(y): f"={c_(y)}{R['mfy']}*{c_(y)}{R['sh']}" for y in F3}}, fmt="#,##0.00")
    sec_("B. 非 AI 與合計")
    row("nai", "非 AI（企業、儲存、邊緣等）", "US$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{p('gn', y)})" for y in F3}}, fmt="#,##0.00")
    ws[f"E{R['sh']}"] = ws[f"E{R['sh']}"].value.replace("{nai}", str(R["nai"]))
    row("rev", "營收", "US$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.00")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI GPU 平台 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("ugpu", "隱含美超微商用 GPU 出貨（AI 營收 ÷ 每顆平均系統價值）", "百萬顆",
        {c_(y): f"={c_(y)}{R['ai']}/((0.5*({c_(y - 1)}{R['uo']}*{c_(y - 1)}{R['vo']}+{c_(y - 1)}{R['un']}*{c_(y - 1)}{R['vn']})+0.5*({c_(y)}{R['uo']}*{c_(y)}{R['vo']}+{c_(y)}{R['un']}*{c_(y)}{R['vn']}))"
                                    f"/(0.5*({c_(y - 1)}{R['uo']}+{c_(y - 1)}{R['un']})+0.5*({c_(y)}{R['uo']}+{c_(y)}{R['un']})))" for y in F3}, fmt="0.00",
        note="對照用：份額 × 會計年度商用 GPU 顆數")
    sec_("C. 損益")
    row("gai0", "AI 毛利率 FY2026 校準值 ＝（毛利 − 非 AI 毛利率 × 非 AI 營收）÷ AI 營收", "%", {"D": f"=(a26_gp-{P['gmn']}*a26_nai)/a26_ai"}, fmt="0.00%",
        note="Derived：FY2026 含 4Q 一次性利益（遞延合約、關稅、存貨跌價回轉），校準值偏高")
    row("gmai", "AI 毛利率", "%", {"D": f"=D{R['gai0']}", "E": f"=$D{R['gai0']}+{p('d27')}", **{c_(y): f"=$D{R['gai0']}+{p('d28')}" for y in (2028, 2029)}}, fmt="0.00%",
        note="A 最高、C 最低（In 頁 C 節）")
    row("gp", "毛利", "US$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"={c_(y)}{R['gmai']}*{c_(y)}{R['ai']}+{P['gmn']}*{c_(y)}{R['nai']}" for y in F3}}, fmt="#,##0.00")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,{P['phi']}*{c_(y)}{R['g']}))" for y in F3}}, fmt="#,##0.00")
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in F3}}, fmt="#,##0.00")
    row("oi", "營業利益", "US$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.00")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    row("nio", "淨利（特別股股利前）÷ 營業利益", "%", {"C": "=a25_ni/a25_oi", "D": "=a26_ni/a26_oi", **{c_(y): f"={P['nio']}" for y in F3}}, fmt="0.0%")
    row("ni0", "淨利（特別股股利前）", "US$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"={c_(y)}{R['oi']}*{c_(y)}{R['nio']}" for y in F3}}, fmt="#,##0.00")
    row("pd", "特別股股利", "US$B", {"C": "=0", "D": "=0", **{c_(y): "=pref_div" for y in F3}}, fmt="#,##0.000",
        note="FY2026 只有約 3 週，已含在實績；2029-06 強制轉換後為 0")
    row("ni", "歸屬母公司淨利", "US$B", {c_(y): f"={c_(y)}{R['ni0']}-{c_(y)}{R['pd']}" for y in FYS}, fmt="#,##0.00",
        note="扣除特別股股利後、歸屬普通股（GAAP）")
    row("nim0", "淨利率（特別股股利前）", "%", {c_(y): f"={c_(y)}{R['ni0']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%", note="收斂法的起點（FY2030 起特別股已轉換）")
    row("sh_", "股數（GAAP 稀釋）", "十億股", {"C": "=a25_sh", "D": "=a26_sh", "E": "=sh27", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in (2028, 2029)}}, fmt="0.0000")
    row("eps", "EPS（GAAP 稀釋）", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("nim", "淨利率（歸屬母公司淨利 ÷ 營收）", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.00%")
    # FY2030–32：收斂法（上游 CY2029 為最後一年，FY2030 需 CY2030；非驅動因子推導）
    for y in FADE:
        c, pc, k = FC[y], FC[y - 1], y - 2029
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+({p('fg')}-$G{R['g']})*{k}/3"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim0']}+({p('fm')}-$G{R['nim0']})*{k}/3"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['sh_']}"] = f"={pc}{R['sh_']}*(1+sh_g)" + ("+conv_sh" if y == 2030 else "")
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh_']}"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("g", "0.0%"), ("rev", "#,##0.00"), ("nim", "0.00%"), ("ni", "#,##0.00"), ("sh_", "0.0000"), ("eps", "0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fm; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 11, "FY2030 起為收斂法：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 特別股股利前淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導；FY2030 股數加特別股轉換股數")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 美超微損益（十億美元；EPS 美元、GAAP 稀釋；歸屬母公司＝扣除特別股股利後）。FY2025A、FY2026A 實績（FY2026 為錨點），FY2027 起為模型。會計年度 7 月～6 月：FY(t) 市場＝½ CY(t−1)＋½ CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數係數 1.0；FY2027 營收＝指引上緣；份額維持；AI 毛利率 +1.5pt",
         "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85；FY2027 營收＝指引中點；份額 ×0.95／×0.92；AI 毛利率 +0.5pt／+0.3pt",
         "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6；FY2027 營收＝指引下緣；份額 ×0.85／×0.80；AI 毛利率 −0.5pt／−1.5pt"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026A（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "US$B", "#,##0.00"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI GPU 平台營收", "US$B", "#,##0.00"),
        ("nai", "  非 AI（企業、儲存、邊緣等）", "US$B", "#,##0.00"), ("aish", "  AI GPU 平台 ÷ 營收", "%", "0.0%"),
        ("gm", "毛利率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"), ("ni", "歸屬母公司淨利", "US$B", "#,##0.00"), ("nim", "  淨利率", "%", "0.00%"),
        ("eps", "EPS", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{PB['A']['nim0']}")
IN.cell(INROW["fade_mB"], 5, f"=(PL!G{PB['B']['nim0']}+nim10)/2")
for k in ("fade_mA", "fade_mB"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 實績").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（US$B）", "+".join(f"ABS(D{PB[s]['rev']}-a26_rev)" for s in SC)),
               ("  三情境：EPS 差（US$）", "+".join(f"ABS(D{PB[s]['eps']}-a26_eps)" for s in SC)),
               ("  情境 A：毛利差（US$B；以 FY2027 公式套用 FY2026）", f"D{PB['A']['gai0']}*D{PB['A']['ai']}+gm_nai*D{PB['A']['nai']}-a26_gp")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 60, sr
PL.cell(2, 1, "灰底＝FY2030–FY2032 收斂法（上游只到 CY2029，FY2030 需 CY2030）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 H 節），非驅動因子推導")
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["K"].width = 70
for c in "CDEFGHIJ": PL.column_dimensions[c].width = 13

# ---------------- Units ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 商用 GPU 系統市場、美超微份額與隱含出貨（情境 A）。前五列為日曆年（CY），其後為美超微會計年度（FY，7 月～6 月）。顆數讀台積電模型、每顆系統價值讀橋接表（唯讀）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "2026", "2027", "2028", "2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("uo", "CY：GB300 世代商用 GPU 顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['uo']}", "0.00", "tsmc 成品"),
        ("un", "CY：Rubin 世代商用 GPU 顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['un']}", "0.00", "tsmc 成品"),
        ("ut", "CY：全部 AI 加速器顆數（對照）", "百萬顆", lambda y: f"=ut_{y % 100}", "0.00", "含 TPU、Trainium、其他 ASIC"),
        ("vn", "CY：每顆系統價值（Rubin 世代）", "千美元/顆", lambda y: f"=PL!{FC[y]}{A_['vn']}", "0.0", "GB300 世代為 In 頁 v_gb"),
        ("mcy", "CY：商用 GPU 系統市場", "US$B", lambda y: f"=PL!{FC[y]}{A_['mcy']}", "#,##0.0", "顆數 × 每顆系統價值"),
        ("mfy", "FY：商用 GPU 系統市場（½ 前一 CY＋½ 當年 CY）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mfy']}" if y >= 2027 else None, "#,##0.0", "FY2026 需 CY2025（上游無），不列"),
        ("sh", "FY：美超微有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}" if y >= 2027 else f"=PL!D{A_['sh0']}", "0.00%", "FY2026 欄＝2H FY26 校準值；FY2027 起以指引校準"),
        ("ai", "FY：美超微 AI GPU 平台營收", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}", "#,##0.0", ""),
        ("ugpu", "FY：隱含美超微商用 GPU 出貨", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['ugpu']}" if y >= 2027 else None, "0.00", "AI 營收 ÷ 每顆平均系統價值"),
        ("gmai", "FY：AI 毛利率", "%", lambda y: f"=PL!{FC[y]}{A_['gmai']}", "0.00%", "FY2026 為校準值"),
        ("gm", "FY：毛利率", "%", lambda y: f"=PL!{FC[y]}{A_['gm']}", "0.00%", "")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in CYS:
        v = f(y)
        if v is None: continue
        UN[f"{cyc(y)}{ur}"] = v; UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：美超微賣的是含 GPU 的整機與整櫃（品牌商，買賣模式），營收隨商用 GPU 顆數 × 每顆系統價值放大；份額在 2H FY26 約數個百分點，FY2027 指引隱含份額明顯上升（大量新訂單）。FY2028 起份額相對值是 Assumed，大客戶（FY2026 占 28%）轉單是最大下檔。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 60; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元，GAAP EPS）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」；目標價＝DCF 與 Comparable 加權（第 0 節）；十年 P/E 規則值（第 3 節，會計年度年底落後 P/E）只作參考"
V["A2"] = "循環位置：美超微不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照廣達、緯穎）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 51
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-23 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "0.00"
for k, y, lab in ((3, 2026, "現價 ÷ FY2026 EPS（實績）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 B）≈ 8／12 × FY2027＋4／12 × FY2028"); V[f"B{O + 7}"] = f"=8/12*PL!E{SUMR[('B', 'eps')]}+4/12*PL!F{SUMR[('B', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 B）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（stockanalysis，10-23）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis，10-23；non-GAAP）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 3}"] = "FY2026 三情境相同（實績）"; V[f"E{O + 7}"] = "2026-11～2027-10：FY2027（7 月～6 月）8 個月＋FY2028 4 個月"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型營收最高的會計年度）").font = BOLD
hdr(V, CY0 + 1, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（美元；約 FY2027 年底＝2027-06，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0.00"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；2027-06 的價格看下一會計年度 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；美超微的低倍數同時反映治理、會計與營運資金風險（見 README）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
# 十年 P/E：會計年度年底（6 月底）落後 P/E；Macrotrends（GAAP TTM EPS）與 fullratio（年度值與 6 月季度 TTM 相同）
PE_MT = [16.35, 18.40, None, None, 26.05, 16.75, 7.63, 21.77, 41.13, 29.52]
PE_FR = [16.57, 17.61, 26.28, 13.82, 16.7, 15.99, 7.34, 20.6, 39.58, 27.69]
T10 = vr + 2
V.cell(T10, 1, "5. 十年 P/E（FY2016–FY2025 會計年度年底＝6 月底；兩來源平均，分位排名用；Macrotrends 2018–2019 因延遲申報缺值）").font = BOLD
hdr(V, T10 + 1, ["年底", "兩來源平均 P/E", "Macrotrends", "fullratio"])
RB0 = T10 + 16
for i, y in enumerate(range(2016, 2026)):
    rr = T10 + 2 + i
    V.cell(rr, 1, str(y))
    V.cell(rr, 2, f"=AVERAGE(C{rr}:D{rr})").number_format = "0.00"
    V.cell(rr, 3, f'=IF(ISNUMBER(B{RB0 + 5 + i}),B{RB0 + 5 + i},"")').number_format = "0.00"
    V.cell(rr, 4, f"=C{RB0 + 5 + i}").number_format = "0.00"
PE10 = f"$B${T10 + 2}:$B${T10 + 11}"
V.cell(T10 + 12, 1, "年數"); V[f"B{T10 + 12}"] = f"=COUNT({PE10})"
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "Macrotrends（6 月底 P/E）", "vals": PE_MT, "tag": "Interested-party（資料商；GAAP 近 12 月 EPS；2018–2019 延遲申報期間 EPS 記 0，缺值）"},
    {"label": "fullratio（年度 P/E）", "vals": PE_FR, "tag": "Interested-party（資料商；年度值與 6 月季度相同，推定為會計年度年底；EPS 口徑可能為調整後）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。美超微會計年度 6 月底結束，取會計年度年底。兩來源 2020 年相差約 9 個倍數點（Macrotrends 的 FY2020 EPS 受重編影響）；FY2024 因股價高峰達約 40x。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025", (V.cell(RB0 + 5, 1).value, V.cell(RB0 + 14, 1).value)
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# 0. 目標價（Andy 2026-10-25 11:28／11:53：DCF 與 Comparable 各半；列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期）
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境（GAAP 稀釋）；FY2030 起為收斂法"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "0.00"); EPSR[y] = 8 + i
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
CFR = {}
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)" for c in "BCD"], "0.00"); CFR[y] = 16 + i
def dcf_formula(c, r="tp_r", g="tp_g", reinv="$B$15"):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"MAX(0,{pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], "0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底、t＝1 對 FY2027（6 月底結束，實際約半年後，略保守）；低於 0 取 0")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("Dell Technologies", "DELL", 16.22, "2026-10-23"), ("Hewlett Packard Enterprise", "HPE", 12.19, "2026-10-23"),
         ("緯穎（Wiwynn）", "6669.TW", 11.73, "2026-10-21"), ("廣達（Quanta）", "2382.TW", 11.52, "2026-10-21"), ("緯創（Wistron）", "3231.TW", 9.21, "2026-10-19")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "AI 伺服器品牌與 ODM 同業：Dell、HPE、緯穎、廣達、緯創（不含 A 股）")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(8/12*B{EPSR[2028]}+4/12*B{EPSR[2029]})+tp_wB*(8/12*C{EPSR[2028]}+4/12*C{EPSR[2029]})+tp_wC*(8/12*D{EPSR[2028]}+4/12*D{EPSR[2029]})", None, None], "0.00",
     "會計年度 6 月底結束：約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 8／12 × FY2028＋4／12 × FY2029")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "", "", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁（0.5）")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "0.00", "stockanalysis（S&P Global）：18 位分析師平均 43.27（中位數 45、最低 15、最高 60），2026-10-22；tickflow 25 位 41.44")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-23 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS（GAAP）÷ 共識前瞻 EPS（non-GAAP，＝價格 ÷ 前瞻 P/E）− 1；期間約晚一年、口徑不同（GAAP 較低）")
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
tset(46, "band", "價格帶（DCF 不如預期～順利）", ['=TEXT(D23,"0.0")&"～"&TEXT(B23,"0.0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 FY2016–FY2025 年底 P/E（第 5 節兩來源平均）的分位；低於歷史最小值記 0%、高於最大值記 100%；只作檢查，不進目標價")
V.cell(48, 1, "目標價方法")
V["B48"] = ("DCF：股東現金流＝GAAP 稀釋 EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為近三年平均約 23.5%（再投資率約 17%），FY2027–32 明確期＋永續；"
            "組裝與品牌商資本支出低（FY2026 1.62 億美元），未改用淨利＋折舊－資本支出；但營運資金占用大（FY2026 營業現金流 −68 億美元、以增資與特別股 56 億補足），以 EPS 推現金流會高估股東實拿的現金，偏誤向上。"
            "Comparable：Dell、HPE、緯穎、廣達、緯創前瞻 P/E 中位數（不含 A 股）；同業倍數未反映美超微的治理與會計折價（自身前瞻 P/E 約 7x），偏誤向上。"
            "FY2030 起為收斂法；FY2030 股數加特別股轉換股數。治理風險不改折現率，以 Sens 案例呈現。"
            "情境對應：A＝順利（顆數基準）、B＝平平（顆數 ×0.8／0.85）、C＝不如預期（顆數 ×0.6）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；FY2027 份額在每個案例內依指引重新校準，部署係數與每顆價值除外）：對 FY2027、FY2028 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("AI 部署 ×0.8", "CY2027 起商用 GPU 顆數 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("FY2027 營收＝指引下緣", "650 億美元（情境 A 為上緣 720 億）", dict(g27="g27_lo")),
         ("份額基準（B 值）", "份額相對值 FY2028 0.95、FY2029 0.92", dict(rel="rel_{y}")),
         ("份額流失（C 值）", "份額相對值 0.85／0.80", dict(rel="rel_{y}_lo")),
         ("治理事件：大客戶流失", "FY2027 指引下緣、FY2028 起份額 ×0.8（出口管制、會計或內控事件使客戶轉向 Dell、ODM）", dict(g27="g27_lo", rel_mult="0.8")),
         ("每顆系統價值 低", "2028–29 Rubin 世代 ×0.85", dict(vpg="vpg_{y}_lo")), ("每顆系統價值 高", "×1.2", dict(vpg="vpg_{y}_hi")),
         ("AI 毛利率 基準（B 值）", "FY2027 +0.5pt、FY2028 起 +0.3pt", dict(d27="dgm_27", d28="dgm_28")),
         ("AI 毛利率 低（C 值）", "FY2027 −0.5pt、FY2028 起 −1.5pt", dict(d27="dgm_27_lo", d28="dgm_28_lo")),
         ("非 AI 毛利率 13%", "基準 15%", dict(gmn="gm_nai_lo")), ("非 AI 毛利率 18%", "基準 15%", dict(gmn="gm_nai_hi")),
         ("非 AI 成長 低", "FY2027 0%、FY2028–29 −5%", dict(gn="g_nai{y}_lo")), ("非 AI 成長 高", "FY2027 +30%、FY2028 +20%、FY2029 +15%", dict(gn="g_nai{y}_hi")),
         ("營業費用率 +0.2pt", "營業費用額外 +0.2% × 營收", dict(opxa="opx_adj_hi")), ("營業費用彈性 0.6", "基準 0.4", dict(phi="phi_hi")),
         ("淨利 ÷ 營業利益 0.68", "利息（營運資金借款）、合規與法律費用上升；1Q FY27 指引隱含值", dict(nio="nio_f_lo")),
         ("淨利 ÷ 營業利益 0.80", "FY2026 實測約 0.805", dict(nio="nio_f_hi"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "gm", 2027, "0.00%"), ("E", "eps", 2027, "0.00"),
                           ("G", "rev", 2028, "#,##0.0"), ("H", "gm", 2028, "0.00%"), ("I", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"F{i}"] = f"=E{i}/E$4-1"; SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c in "FJ": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 B", "CY2028 顆數 ×0.8、CY2029 ×0.85；指引中點；份額、毛利率 B 值", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "CY2028–29 顆數 ×0.6；指引下緣；份額、毛利率 C 值", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 營收以指引校準，所以 FY2027 主要受指引區間、毛利率與淨利係數影響；FY2028 起顆數情境、份額與每顆價值才發揮作用。美超微淨利率只有約 5%–6%，AI 毛利率每 1 個百分點約影響 EPS 10%。治理事件以份額與淨利係數案例呈現，不改情境結構。")
SENS_NOTE = nn + 3
TPS0 = nn + 5
SE.cell(TPS0, 1, "目標價參數（只改 DCF 參數；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "#,##0.00"
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.3", "", dict(w="0.3")), ("DCF 權重 0.7", "", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
assert TPS_LAST < 40, TPS_LAST
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 48
for c in "CDEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引、共識與份額口徑。只作事後對照，不為對上數字而調參數。共識 EPS 為 non-GAAP；本模型為 GAAP，另列「模型＋股票報酬調整」對照"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 共識", "=c_rev27", pl("rev", 2027), "stockanalysis／Yahoo（15 位）", "本模型 FY2027 營收＝公司指引（A 上緣、B 中點、C 下緣）")
rec("FY2028 營收 vs 共識", "=c_rev28", pl("rev", 2028), "Yahoo（15 位）", "")
rec("FY2027 EPS（GAAP）vs 共識（non-GAAP）", "=c_eps27", pl("eps", 2027), "stockanalysis／Yahoo", "口徑不同：GAAP 較 non-GAAP 低約股票報酬")
rec("FY2027 EPS（模型＋股票報酬調整）vs 共識", "=c_eps27", lambda s: f"=PL!E{PB[s]['eps']}+sbc_ps", "同上", "調整＝每季股票報酬（稅後）× 4 ÷ 股數（In 頁 sbc_ps）")
rec("FY2028 EPS（模型＋股票報酬調整）vs 共識", "=c_eps28", lambda s: f"=PL!F{PB[s]['eps']}+sbc_ps", "Yahoo（17 位）", "")
RC.cell(rr_, 1, "B. 公司指引與初步數字").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 指引中點（Interested-party）", "=g27", pl("rev", 2027), "公司 2026-08-11", "依情境取指引上緣／中點／下緣，差距即指引區間")
rec("FY2027 毛利率 vs 1Q FY27 初步毛利率（百分點）", "=q1_gm", pl("gm", 2027), "公司 2026-10-20", "1Q 含 4Q FY26 遞延合約的組合利益", fmt="0.00%", pt=True)
rec("1Q FY27 初步營收 × 4 vs FY2027 營收", "=q1_rev*4", pl("rev", 2027), "公司 2026-10-20", "1Q 年化已高於指引上緣；全年指引尚未調整（11-03 財報）")
RC.cell(rr_, 1, "C. 份額與市場").font = BOLD; rr_ += 1
rec("2H FY26 有效份額（商用 GPU 系統市場）", None, lambda s: f"=PL!D{PB[s]['sh0']}", "Derived", "", fmt="0.00%")
rec("FY2027 有效份額（指引校準）", None, pl("sh", 2027), "Derived", "", fmt="0.00%")
rec("FY2027 份額 ÷ 2H FY26 份額", None, pl("shr", 2027), "Derived", "大於 1：指引隱含份額上升（4Q FY26 逾 600 億美元新訂單）", fmt="0.00")
rec("FY2027 隱含商用 GPU 出貨（百萬顆）", None, pl("ugpu", 2027), "Derived", "", fmt="0.00")
RC.cell(rr_, 1, "D. 校準").font = BOLD; rr_ += 1
rec("AI 毛利率 FY2026 校準值（Derived）", None, lambda s: f"=PL!D{PB[s]['gai0']}", "FY2026 實績", "含 4Q 一次性利益；不含 4Q 的 FY2026 毛利率見 Q 頁 B24", fmt="0.00%")
RC.cell(rr_, 1, "E. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代「台積電顆數 × 情境係數」"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文為新聞稿與財報說明稿；10-K 內容只經二手摘要，原文待人工複核；券商與資料商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：FY2026 第四季與全年新聞稿（2026-08-11；營收、毛利率、營業費用、淨利、EPS、股數、1Q FY27 與 FY2027 指引、董事會調查）", PR4Q),
        ("公司：4Q FY26 財報說明稿（AI 占比、客戶與地區結構、最大客戶 28%、毛利率說明、新訂單）", PREP),
        ("公司（摘要）：MarketBeat 2026-08-11 法說重點", MB0811),
        ("公司：2Q FY26 新聞稿（8-K；1Q、2Q FY26 數字）", PRQ2),
        ("公司：1Q FY27 初步營運更新（2026-10-20；營收 185～195 億、毛利率 12.0–12.5%、11-03 財報）", PRQ1U),
        ("公司：2026-06-11 普通股與強制可轉換特別股定價（股利率 7%、轉換率、2029-06-01 轉換）", FIN0611),
        ("公司：2026-08 獨立調查完成公告（出口管制）", INVEST),
        ("媒體：Stocktwits 2026-08-20（調查結論、2026-03 起訴、共同創辦人辭任董事）", ST0820),
        ("公司：FY2026 10-K（2026-08-31）", K10),
        ("媒體：StockTi（10-K 摘要：內控仍未有效、BDO 否定意見、單一供應商 63.1%）", STK10),
        ("媒體：Futurum（2Q FY26 平台）", FUT),
        ("資料商：stockanalysis 統計頁（現價、股數、P/E、目標價）", SA),
        ("資料商：stockanalysis 預測頁（共識營收、EPS、目標價）", SAF),
        ("資料商：stockanalysis 比率頁（ROE、年度 P/E）", SAR),
        ("資料商：stockanalysis 損益表（FY2022–FY2026）", SAI),
        ("資料商：Yahoo Finance 分析頁（FY2027–28 共識）", YF),
        ("資料商：tickflow 預測頁（目標價第二來源）", TICK),
        ("資料商：Simply Wall St（目標價歷史、ROE 預測、1Q 指引上修）", SWS),
        ("資料商：Macrotrends P/E（6 月底）", MT), ("資料商：Macrotrends 營收", MTREV), ("資料商：Macrotrends 淨利", MTNI),
        ("資料商：fullratio P/E（年度、現價）", FR),
        ("同業前瞻 P/E：stockanalysis（Dell）", P_DELL), ("同業前瞻 P/E：stockanalysis（HPE）", P_HPE), ("同業前瞻 P/E：stockanalysis（緯穎）", P_WY),
        ("同業前瞻 P/E：stockanalysis（廣達）", P_QC), ("同業前瞻 P/E：stockanalysis（緯創）", P_WS),
        ("市場層：tsmc、markets/chip_bridge 成品（本 repo，唯讀）", f"{rel(TSMC_F)}；{rel(BR_F)}"),
        ("待補（人工）：1Q FY27 財報（2026-11-03）；10-K 原文（客戶占比、內控）；各季 AI GPU 占比的公司原文；AI 與非 AI 毛利率；可轉債轉換價與稀釋；特別股超額配售執行數；前瞻 P/E 歷史（Bloomberg）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("美超微 Super Micro Computer（SMCI US）收支與估值模型 v0.1.1（AI 半導體 Project：伺服器組裝次產業，樣板＝wiwynn/、quanta/）　2026-10-25", True),
 ("版本：v0.1.1（2026-10-25：補上檔空間級距、投資建議、重點觀察（數值不變））。v0.1（初版）。美超微是 AI 伺服器品牌商（整機、整櫃、直接液冷），客戶含 neocloud、大型資料中心與企業，以買賣模式銷售含 NVIDIA、AMD GPU 的系統。商用 GPU 顆數與情境係數讀台積電模型，每顆系統價值讀橋接表（Tokenomics 整櫃價格），本模型只決定美超微的份額、毛利率與非 AI 營收。", False),
 ("模型命題：美超微 AI 營收＝商用 GPU 顆數（台積電模型，同字母情境）× 每顆系統價值（整櫃買方價 ÷ 每架 GPU）× 美超微有效份額（FY2027 以公司指引校準，FY2028 起依情境變動）；非 AI 營收以成長率推；毛利率＝AI 與非 AI 加權，AI 毛利率依情境調整。", False),
 ("核心問題：FY2027 營收指引 650～720 億美元（1Q 已初報約 190 億）能否轉成 GAAP EPS 約 5 美元、並守住約 11% 毛利率；現價約 33 美元、約 7 倍前瞻 P/E 的治理折價是否過深。", False), ("", False),
 ("會計年度：FY(t)＝t−1 年 7 月～t 年 6 月；FY2026 已結束（實績錨點）。上游為日曆年：FY(t) 市場＝½ CY(t−1)＋½ CY(t)。金額：十億美元；EPS：GAAP 稀釋（美元）。口徑：歸屬母公司淨利＝扣除強制可轉換特別股股利後、歸屬普通股；2029-06 特別股轉換後（FY2030 起）股利消失、股數增加。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 營收：tsmc 成品的商用 GPU 顆數（Blackwell、Hopper、Rubin、Rubin Ultra、AMD）× 情境顆數係數 × 每顆系統價值（橋接表：GB300 機架 430 萬 ÷ 72、VR200 機架 840 萬 ÷ 72 美元）＝商用 GPU 系統市場；× 美超微有效份額（FY2027＝（指引營收 − 非 AI）÷ 基準市場，Derived；FY2028 起 × 份額相對值，Assumed）。", False),
 ("非 AI：FY2026 殘差 × 成長率（Assumed）。毛利率：非 AI 15%（Assumed）；AI 毛利率＝FY2026 校準值＋情境調整（Interested-party／Assumed）。營業費用彈性 phi（Derived）。淨利：營業利益 × 淨利係數（稅與利息）− 特別股股利。", False),
 ("校準：FY2026 ＝ 全年實績（公司 2026-08-11）；FY2027 營收＝公司指引（A 上緣、B 中點、C 下緣）。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2）＋0.5 × Comparable 目標價（Dell、HPE、緯穎、廣達、緯創前瞻 P/E 中位數 × 12 個月後前瞻 EPS（8／12 × FY2028＋4／12 × FY2029）；不含 A 股）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解。十年 P/E 規則值 × FY2027 EPS 保留為參考（不是目標價）。情境對應：A＝順利、B＝平平、C＝不如預期（原為顆數情境）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 特別股股利前淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文為新聞稿與財報說明稿；10-K 只經二手摘要；券商與資料商為 Interested-party。", False),
 ("", False),
 ("已知限制：(1) 治理風險：2024 年會計師辭任與延遲申報、2026-03 前員工與共同創辦人因出口管制被起訴、FY2026 10-K 內控仍未有效（BDO 否定意見）；不改情境結構，以 Sens 案例（大客戶流失、淨利係數）呈現；(2) FY2027 份額以公司指引校準，FY2028 起份額、AI 毛利率為 Assumed；(3) 最大客戶占 FY2026 營收 28%、單一供應商占採購 63.1%；(4) 每顆系統價值以 NVIDIA 整櫃買方價近似，HGX 與 AMD 為 Analogy；(5) EPS 推現金流未扣營運資金（FY2026 營業現金流 −68 億美元），DCF 偏高；(6) 共識 EPS 為 non-GAAP，本模型為 GAAP；(7) 1Q FY27 財報（2026-11-03）後更新；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ================= v0.1.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（工作單；Andy 2026-10-25 14:00、14:17） =================
# 全部放在 Val、README 現有最後一列之後，不位移既有儲存格
OBS = [
       "目標價 US$76.6（DCF 78.0、Comparable 75.2 各半），比現價 32.94（2026-10-23）高 132.5%，級距大於 50%；Comparable 單獨為 128.3%。主因是倍數與 EPS 兩者：同業（Dell、HPE、緯穎、廣達、緯創）前瞻 P/E 中位數 11.7 倍，現價隱含只有 7.1 倍、共識目標價 43.27 隱含 9.3 倍（比同業低 21%），Comparable 未計美超微的治理與會計折價；本模型 12 個月後前瞻 EPS 6.41（GAAP）也比共識前瞻 EPS 4.67（non-GAAP，期間約早一年）高 37.4%。",
       "AI 毛利率是最敏感的變數：取 C 值（FY2027 −0.5pt、FY2028 起 −1.5pt）時 FY2028 EPS −30.1%，取 B 值時 −12.0%；1Q FY27 初步毛利率 12.25% 含遞延合約的組合利益。1Q FY27 財報（2026-11-03）的全年指引是第一個檢驗點。",
       "情境分歧大：FY2028 EPS A／B／C＝8.60／6.07／3.58（C 比 A 低 58.4%），DCF 情境價值 128.1／66.5／31.4；治理事件（大客戶流失）案例 FY2028 EPS −28.8%，價格帶下緣 31.4 已低於現價。"
      ]   # 重點觀察（與 README.md 相同文字；數字取自重算後的 xlsx）
_U = f"B{T['upw']}"
GR = V.max_row + 3
V.cell(GR - 1, 1, "0.4 上檔空間級距與投資建議（v0.1.1 新增；放在本頁最後以免位移既有儲存格；依「上檔空間：加權目標價」第 " + str(T["upw"]) + " 列）").font = BOLD
tset(GR, "grade", "上檔空間級距", [f'=IF({_U}>0.5,"大於 50%",IF({_U}>0.25,"25%–50%",IF({_U}>=0,"0%–25%",IF({_U}>=-0.25,"0%～−25%",IF({_U}>=-0.5,"−25%～−50%","小於 −50%")))))', None, None], "General",
     "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）；負值區間同樣重要")
tset(GR + 1, "rec", "投資建議", [f'=IF({_U}>0.5,"強力買進",IF({_U}>0.25,"買進",IF({_U}>=-0.25,"中立",IF({_U}>=-0.5,"放空","強力放空"))))', None, None], "General",
     "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）")
rn = RD.max_row + 1
RD.cell(rn, 1, f'="投資建議："&Val!B{GR + 1}&"（依上檔空間固定規則）"').font = BOLD
OBS_ROW = rn + 1
for i_, t_ in enumerate(OBS, 1):
    RD.cell(rn + i_, 1, f"重點觀察 {i_}：" + t_)


# ---------------- 分部原則第 3 條（2026-10-28）：核對列與宣告（追加於工作表最後，不位移既有列，數值不變） ----------------
import os as _os2, sys as _sys2
_sys2.path.insert(0, _os2.path.join(_os2.path.dirname(_os2.path.dirname(_os2.path.abspath(__file__))), "tools", "segment_check"))
from addrow import add_check
add_check(Q, seg_rows=[16, 17], total_row=4, cols="BCDEF", rnd=3, note="分部原則第 3 條（2026-10-28 追加）；分部宣告見 segments.json")
RD.cell(RD.max_row + 1, 1, "v0.1.2（2026-10-28）：分部原則重建（數值不變）。Q 頁最後追加核對列「檢查：分部合計 − 合併營收」；官方分部與子列宣告見 segments.json。")
if isinstance(RD.cell(1, 1).value, str): RD.cell(1, 1).value = RD.cell(1, 1).value.replace("v0.1.1", "v0.1.2")

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# 相對倍數調整（Andy 2026-10-27「照建議做」；tools/relative_multiple.py，2026-10-29 接入）
import os, sys  # noqa: E401
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tools"))
import relative_multiple as RMX  # noqa: E402
RMX.apply(wb, "smci")
wb.save(sys.argv[1])
if os.environ.get("SMCI_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE, "UN_NOTE": UN_NOTE,
               "T": T, "T10": T10, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "PEER_LAST": PEER_LAST, "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0, "NCASES": len(cases)},
              open(os.environ["SMCI_ROWMAP"], "w"), ensure_ascii=False, indent=1)
