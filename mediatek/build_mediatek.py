# 聯發科（2454.TW）收支與估值模型 v0.1.1（AI 半導體 Project：運算層，手機 SoC＋智慧邊緣平台＋電源 IC＋AI ASIC；ASIC 樣板 avgo/、結構參考 mrvl/、amkr/）
# 用法：python3 mediatek/build_mediatek.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；ASIC 每顆價格：千美元（US$K）；EPS：新台幣（基本）。
# 口徑：歸屬母公司淨利（TIFRS）。AI ASIC 營收＝Google TPU（與其他 ASIC）顆數 × 聯發科份額 × 每顆營收（avgo 規格）× k_x × 匯率。
import sys, os, glob
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
XP = ("tpu", "oth")
PNAME = {"tpu": "Google TPU", "oth": "其他 ASIC（MTIA、Maia、Ascend 等）"}

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); AVGO_F = only("avgo/dist/*.xlsx"); MRVL_F = only("mrvl/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith(("u_", "g29_", "um"))):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
for k in XP:
    assert all(f"u_{k}_{y}" in TS for y in (26, 27, 28)) and f"g29_{k}" in TS, k
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC)
AV = {}
for r in openpyxl.load_workbook(AVGO_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith(("core_", "sh_", "gb_", "hbm_", "u_")) or k in ("f_hbm", "m_hbm")):
        AV[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(AVGO_F)}｜In!E{r[0].row}")
for k in XP:
    for y in (26, 27, 28):
        assert abs(AV[f"u_{k}_{y}"][0][1] - TS[f"u_{k}_{y}"][0]) < 1e-9, (k, y)   # avgo 與 tsmc 同一份顆數
MV = {}
for r in openpyxl.load_workbook(MRVL_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith(("msh_tpu", "msh_oth")):
        MV[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(MRVL_F)}｜In!E{r[0].row}")
assert len(MV) == 6, MV.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B）；ASIC 每顆價格：千美元（US$K）；顆數：百萬顆；匯率：新台幣／美元。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司原文＝聯發科法說營運報告、新聞稿、法說逐字稿與講稿（mediatek.com）；公司指引與目標為 Interested-party；券商、資料商、媒體為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
MTK = "https://www.mediatek.com/hubfs/MediaTek%20Assets/Pdfs/Quarterly%20Earnings%20Release/"
R2Q = MTK + "2026/Quarterly%20Earnings%20Release-2026Q2/%5B%E6%B3%95%E8%AA%AA%E6%9C%83%E8%B3%87%E6%96%99%5DQ2%E7%87%9F%E9%81%8B%E5%A0%B1%E5%91%8A.pdf"
P2Q = MTK + "2026/Quarterly%20Earnings%20Release-2026Q2/Prepared%20remarks.pdf"
T1Q = MTK + "2026/Quarterly%20Earnings%20Release-2026Q1/%E8%AA%AA%E6%98%8E%E6%9C%83%E9%80%90%E5%AD%97%E7%A8%BF.pdf"
R4Q = MTK + "2025/Quarterly%20Earnings%20Release-2025Q4/%5B%E6%B3%95%E8%AA%AA%E6%9C%83%E8%B3%87%E6%96%99%5DQ4%E7%87%9F%E9%81%8B%E5%A0%B1%E5%91%8A.pdf"
T4Q = MTK + "2025/Quarterly%20Earnings%20Release-2025Q4/%E8%AA%AA%E6%98%8E%E6%9C%83%E9%80%90%E5%AD%97%E7%A8%BF.pdf"
SEP26 = "https://finance.technews.tw/2026/10/08/mtk-2454-202609-financial-report/"
SDOG = "https://statementdog.com/news/17861"
STORM = "https://www.storm.mg/article/11153222"
SA = "https://stockanalysis.com/quote/tpe/2454/"
SAF = "https://stockanalysis.com/quote/tpe/2454/forecast/"
KGI = "https://www.kgi.com.hk/zh-cn/-/media/files/kgishk/research-reports/tw-reports/mediatek-2454-tt_20260723_en.pdf"
UDN8 = "https://udn.com/news/story/7252/9668107"
UDN5 = "https://money.udn.com/money/amp/story/5607/9479348"
URLNAME = {R2Q: "公司 2Q26 營運報告", P2Q: "公司 2Q26 法說講稿", T1Q: "公司 1Q26 法說逐字稿", R4Q: "公司 4Q25 營運報告", T4Q: "公司 4Q25 法說逐字稿",
           SEP26: "科技新報（9 月營收，2026-10-08）", SDOG: "財報狗（3Q26 營收，2026-10）", STORM: "風傳媒（2Q26 法說，2026-07-31）",
           SA: "stockanalysis.com", SAF: "stockanalysis.com 共識頁", KGI: "凱基（KGI）2026-07-23 報告", UDN8: "經濟日報 2026-08-03", UDN5: "經濟日報 2026-05-03"}

sec("A. AI 加速器顆數（tsmc 成品；與 HBM、台積電、Broadcom、Marvell 模型同一份；情境係數讀台積電模型，不另訂）")
for k in XP:
    for i, y in enumerate((26, 27, 28)):
        inp(f"u_{k}_{y}", f"{PNAME[k]} 顆數 CY20{y}", "百萬顆", None, TS[f"u_{k}_{y}"][0], None, "Interested-party" if k == "tpu" else "Derived／Assumed",
            TS[f"u_{k}_{y}"][1], "avgo 成品同值（建置時檢查）" if i == 0 else "同上")
    inp(f"g29_{k}", f"{PNAME[k]} CY2029 顆數成長", "%", None, TS[f"g29_{k}"][0], None, "Assumed", TS[f"g29_{k}"][1], "markets/memory 無 2029；沿用台積電模型")
for s in SC:
    for y in (28, 29):
        inp(f"um{y}_{s}", f"情境 {s}：CY20{y} 顆數 ÷ 基準", "x", None, TS[f"um{y}_{s}"][0], None, "Assumed", TS[f"um{y}_{s}"][1],
            "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、Broadcom、Marvell 模型相同）" if (s == "A" and y == 28) else "同上")

sec("B. ASIC 每顆規格（權威在 Broadcom 模型 In 頁，唯讀；本模型只加聯發科份額）與份額")
for k, lab in (("core_tpu", "TPU 每顆核心售價（不含 HBM）"), ("core_oth", "其他 ASIC 每顆核心售價（不含 HBM）")):
    lo, b_, hi = AV[k][0]
    inp(k, lab, "US$K", lo, b_, hi, "Assumed", AV[k][1], "Broadcom 模型的 ASIC 每顆規格（本專案權威）；本模型不另訂；聯發科負責的世代與 Broadcom 是否同價：未找到")
lo, b_, hi = AV["f_hbm"][0]
inp("f_hbm", "HBM 由 ASIC 設計夥伴代購並計入營收的比例", "x", lo, b_, hi, "Assumed", AV["f_hbm"][1], "沿用 Broadcom 模型；聯發科客戶是否自購 HBM：未找到（1Q26 法說未回答）")
lo, b_, hi = AV["m_hbm"][0]
inp("m_hbm", "代購 HBM 的毛利率", "x", lo, b_, hi, "Assumed", AV["m_hbm"][1], "沿用 Broadcom 模型")
for k in XP:
    for y in (26, 27, 28):
        inp(f"gb_{k}_{y}", f"{PNAME[k]} 每顆 HBM CY20{y}", "GB", None, AV[f"gb_{k}_{y}"][0][1], None, "Derived／Assumed", AV[f"gb_{k}_{y}"][1], "markets/memory UNITS；CY2029 沿用 CY2028" if y == 26 else "同上")
for k in XP:
    for y in (26, 27, 28):
        inp(f"av_{k}_{y}", f"Broadcom 占 {PNAME[k]} 顆數 CY20{y}（份額加總檢查用）", "%", None, AV[f"sh_{k}_{y}"][0][1], None, "Assumed", AV[f"sh_{k}_{y}"][1], "Broadcom 模型；Broadcom＋Marvell＋聯發科（高）≤ 100%（建置時檢查）")
        inp(f"mv_{k}_{y}", f"Marvell 占 {PNAME[k]} 顆數 CY20{y}（份額加總檢查用）", "%", None, MV[f"msh_{k}_{y}"][0][1], None, "Assumed", MV[f"msh_{k}_{y}"][1], "Marvell 模型")
# mtk_* 份額：Siliconomics In 頁 M7（v0.3 起單一來源；數值與文字由本檔原寫死值逐字搬入）
for k in XP:
    for y in (26, 27, 28):
        _r = SIL.V[f"mtk_{k}_{y}"]; lo, b_, hi = _r["lo"], _r["base"], _r["hi"]
        inp(*SIL.row9(f"mtk_{k}_{y}"))
        assert AV[f"sh_{k}_{y}"][0][1] + MV[f"msh_{k}_{y}"][0][1] + hi <= 1.0 + 1e-9, (k, y)   # Broadcom＋Marvell（基準）＋聯發科（高）≤ 100%
inp("pfc", "CY2028 起 ASIC 核心售價係數（第二顆 ASIC；Sens 用；基準 1）", "x", 0.8, 1.0, 2.0, "Assumed", "經濟日報 2026-05-03 轉述高盛：下一代 ASIC 平均單價為現行數倍（Interested-party）；凱基：v9 單價至少 +200%", "基準沿用 Broadcom 模型規格（不另訂）；高值 2.0 只放 Sens")
inp("uf", "聯發科 ASIC 顆數係數（Sens 用；CY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "無", "±20%")
inp("m_core", "ASIC 核心（不含代購 HBM）毛利率", "x", 0.35, 0.45, 0.55, "Assumed", "公司：ASIC 放量將使營業利益率上升（營業費用低；4Q25 法說）；ASIC 毛利率未揭露", "第二來源：未找到")
inp("hf26", "CY2026 TPU 每顆營收中代購 HBM 的比例（Q 頁第四季毛利用）", "x", None, "=f_hbm*gb_tpu_26*hbm_A_26/1000/(core_tpu+f_hbm*gb_tpu_26*hbm_A_26/1000)", None, "Derived", "In 頁 B、C 節", "—")

sec("C. 記憶體價格路徑（讀 Broadcom 模型成品，其來源為 skhynix 成品；唯讀；情境別）")
for s in SC:
    for y in (26, 27, 28, 29):
        inp(f"hbm_{s}_{y}", f"HBM 混合單價 CY20{y}（{s}）", "US$/GB", None, AV[f"hbm_{s}_{y}"][0][1], None, "Derived", AV[f"hbm_{s}_{y}"][1], "markets/memory：2026 BofA $14.6/GB；2027 +80%" if (s == "A" and y == 26) else "同上")

sec("D. 手機、智慧邊緣平台（不含 ASIC）、電源 IC：簡化成長率（FY2027 起；情境別）")
GR = {"mob": ("手機", {"A": (0.05, 0.04), "B": (0.0, 0.02), "C": (-0.08, -0.03)}, (-0.12, 0.08),
              "公司：2026 年全球智慧型手機出貨約 −15%（記憶體成本）、2027 年持續擴大旗艦市占（2Q26 講稿）", "FY2026 錨點手機約 −17%（Q 頁）；C＝記憶體成本續壓、市占流失"),
      "sep": ("智慧邊緣平台（不含 ASIC）", {"A": (0.08, 0.07), "B": (0.04, 0.04), "C": (-0.03, 0.01)}, (-0.06, 0.12),
              "公司：連網、電視 SoC（含 DRAM 內容）、車用、運算市占提升；2Q26 年增 26%（2Q26 講稿）", "FY2026 錨點約 +26%（含記憶體轉嫁，偏高基期）"),
      "pic": ("電源 IC", {"A": (0.08, 0.08), "B": (0.04, 0.04), "C": (0.0, 0.0)}, (-0.05, 0.12),
              "公司：運算與資料中心市占提升（2Q26 講稿）", "第二來源：未找到")}
for g, (nm, d, (lo, hi), s1, s2) in GR.items():
    for s in SC:
        inp(f"g_{g}27_{s}", f"{nm}年增 FY2027（情境 {s}）", "%", lo, d[s][0], hi, "Assumed", s1 if s == "A" else "同上", s2 if s == "A" else "—")
        inp(f"g_{g}L_{s}", f"{nm}年增 FY2028–29（情境 {s}）", "%", lo, d[s][1], hi, "Assumed", "同上", "—")
inp("gmadj", "非 ASIC 毛利率調整（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "非 ASIC 毛利率為 FY2026 殘差（Derived），之後沿用")

sec("E. FY2026 錨點（1H26 實績＋3Q26 營收實績＋3Q26 指引＋4Q26 假設）")
inp("dc26", "公司 FY2026 資料中心（AI ASIC 為主）營收", "US$B", 2.0, 2.0, 2.5, "Interested-party", P2Q + "：超過 20 億美元（2Q26 講稿）", "經濟日報 2026-08-03 轉述高盛約 22 億美元；4Q25 法說為「超過 10 億美元」（一年內上修）")
inp("asic13", "其中 1–3Q26（量產前：NRE、樣品等；Q 頁全數放在 3Q）", "US$B", 0.1, 0.3, 0.6, "Assumed", "公司：第一顆 ASIC 於 4Q26 量產（2Q26 講稿）", "1Q26 逐字稿：「第四季約 20 億美元」說法與年度說法不易區分（原文含糊）")
inp("fx_q1", "1Q26 匯率（公司指引假設）", "NT$/US$", None, 31.2, None, "Interested-party", T4Q + "：1Q26 指引匯率 31.2", "實際平均匯率：未找到")
inp("fx_q2", "2Q26 平均匯率", "NT$/US$", None, 31.6, None, "Verified", P2Q, "風傳媒同值")
inp("fx26", "3Q、4Q26 匯率（公司 3Q26 指引假設；FY2026 ASIC 換算用）", "NT$/US$", 30.5, 32.0, 33.0, "Interested-party", P2Q + "：3Q26 指引匯率 32", "實際平均匯率：未找到")
inp("q3_rev", "3Q26 營收（月營收實績；7–9 月合計）", "NT$B", None, "=469.391-301.334", None, "Verified", SEP26 + "：前 9 月 4,693.91 億；1H26 3,013.34 億（公司）", SDOG + "：3Q26 1,680.58 億，高於財測上緣 1,598 億")
inp("q3_mob", "3Q26 手機季增（公司指引：持平至中個位數下滑）", "%", -0.06, -0.03, 0.0, "Interested-party", P2Q, "只用來拆分 3Q 營收（等比例放大到實績）")
inp("q3_sep", "3Q26 智慧邊緣平台季增（公司指引：中至高個位數）", "%", 0.04, 0.07, 0.09, "Interested-party", P2Q, "同上")
inp("q3_pic", "3Q26 電源 IC 季增（公司指引：約持平）", "%", -0.02, 0.0, 0.02, "Interested-party", P2Q, "同上")
inp("q3_gm", "3Q26 毛利率（指引 46% ± 1.5pt 中點）", "%", 0.445, 0.46, 0.475, "Interested-party", P2Q, "3Q26 財報約 10 月底公布（確切日期未找到）")
inp("q3_opx", "3Q26 營業費用（指引費用率 31% × 營收指引中點）", "NT$B", 46.0, 48.5, 51.0, "Interested-party", P2Q + "：費用率 31% ± 2pt、營收 1,522–1,598 億", "2Q26 47.44（公司）")
inp("q4_mob", "4Q26 手機季增", "%", -0.12, -0.05, 0.03, "Assumed", "公司：旗艦新機帶動 4Q 出貨（財報狗轉述）", "供應鏈：11 月起手機需求轉弱（財報狗，Interested-party）")
inp("q4_sepx", "4Q26 智慧邊緣平台（不含 ASIC）季增", "%", -0.10, -0.05, 0.0, "Assumed", "3Q 為消費性旺季（電視、連網）；4Q 季節性下滑", "—")
inp("q4_pic", "4Q26 電源 IC 季增", "%", -0.08, -0.03, 0.03, "Assumed", "無", "—")
inp("q4_gmx", "4Q26 非 ASIC 毛利率", "%", 0.445, 0.46, 0.47, "Assumed", "公司：全年毛利率維持在 3Q 指引區間（2Q26 講稿）", "1Q26 46.3%、2Q26 46.2%（公司）")
inp("q4_opx", "4Q26 營業費用", "NT$B", 48.0, 50.0, 53.0, "Assumed", "研發費用逐季增加（1Q 38.3、2Q 39.4；公司）", "—")
inp("q_nonop", "3Q、4Q26 營業外淨額（每季）", "NT$B", 3.5, 4.8, 5.5, "Assumed", "1Q26 4.13、2Q26 5.09（公司）", "US$5B 融資額度若動用，利息收入下降（Assumed）")
inp("q_tax", "3Q、4Q26 有效稅率", "%", 0.10, 0.13, 0.16, "Assumed", "1Q26 9.8%、2Q26 12.0%、FY2025 15.0%（公司）", "—")
inp("q_nci", "3Q、4Q26 非控制權益（每季）", "NT$B", None, 0.25, None, "Derived", "1Q26 0.222、2Q26 0.270（公司）", "—")
inp("q_sh", "3Q、4Q26 加權股數", "十億股", None, "=24.335/15.28", None, "Derived", "2Q26 歸屬母公司淨利 ÷ EPS（公司）", "stockanalysis：流通股數 1.60B")

sec("F. FY2027 起：營業費用、營業外、稅、股數、匯率")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.2, 0.3, 0.6, "Assumed", "ASIC 營收放量所需營業費用低（公司 4Q25 法說：營業利益率上升）；FY2025 營業費用 +11.6%、營收 +12.3%（公司）", "高盛 2027 營業利益率 22%、2028 33%（Interested-party，經濟日報）")
inp("g_opx_min", "營業費用年增下限", "%", 0.03, 0.06, 0.10, "Assumed", "研發費用年增 FY2025 +12.4%（公司）", "—")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 +2 個百分點")
inp("nonop_f", "營業外淨額（每年）", "NT$B", 10, 18, 24, "Assumed", "FY2025 21.4、1H26 9.2（公司）", "現金與流動金融資產 2Q26 198.4（公司）；US$5B 融資額度（2Q26 講稿）")
inp("tax", "有效稅率（FY2027 起）", "%", 0.12, 0.14, 0.17, "Assumed", "FY2024 10.4%、FY2025 15.0%、1H26 10.9%（公司）", "—")
inp("nci", "非控制權益（每年）", "NT$B", None, 1.0, None, "Derived", "FY2025 0.80、1H26 0.49（公司）", "—")
inp("sh_g", "加權股數年變動", "%", -0.005, 0.0, 0.005, "Assumed", "FY2025 1.592 → 2Q26 1.593 十億股（公司）", "員工酬勞股票：小")
inp("fx", "匯率（FY2027 起；ASIC 美元營收換算）", "NT$/US$", 30.0, SIL.base("fx_twd"), 34.0, "Assumed", "公司 3Q26 指引假設 32", "非 ASIC 以新台幣成長率推，未另換算")

sec("G. 公司展望、目標與共識、券商（只作 Recon 對帳，不作校準）")
inp("co_usd25", "公司 FY2025 美元營收", "US$B", None, 19.1, None, "Verified", T4Q + "：191 億美元（年增 15.6%）", "經濟日報同")
inp("co_g26", "公司 FY2026 美元營收年增（高個位數、目標區間上緣）", "%", 0.07, 0.09, 0.10, "Interested-party", P2Q, "1Q26 法說為中至高個位數；3Q26 營收已超出指引")
inp("co_sam27", "公司 FY2027 資料中心 ASIC 可服務市場（SAM）", "US$B", 70, 80, 80, "Interested-party", P2Q + "：800 億美元", T1Q + "：700–800 億美元")
inp("co_sh27", "公司 FY2027 SAM 份額目標（15–20% 中點）", "%", 0.15, 0.175, 0.20, "Interested-party", P2Q + "：由 10–15% 上調", T4Q + "：10–15%（「相當安全」）")
inp("co_dc27sh", "公司說法：FY2027 ASIC 占營收約 20%「可能」", "%", None, 0.20, None, "Interested-party", T4Q, "—")
inp("cons_rev26", "共識營收 FY2026", "NT$B", 613.8, 662.3, 711.4, "Interested-party", SAF + "（25 位分析師；2026-10-23 擷取）", "凱基 678.0")
inp("cons_eps26", "共識 EPS FY2026", "NT$", 58.48, 68.27, 73.35, "Interested-party", SAF + "（25 位）", "凱基 68.97、Citi 68.53")
inp("kgi_rev27", "凱基 FY2027 營收", "NT$B", None, 1110.805, None, "Interested-party", KGI, "—")
inp("kgi_eps27", "凱基 FY2027 EPS", "NT$", None, 112.60, None, "Interested-party", KGI, "Citi 133.87（經濟日報 2026-08-03）；高盛 132.18（2026-05）")
inp("citi_eps27", "Citi FY2027 EPS", "NT$", None, 133.87, None, "Interested-party", UDN8, "—")
inp("kgi_eps28", "凱基 FY2028 EPS", "NT$", None, 246.60, None, "Interested-party", KGI, "Citi 260.38；高盛 406.51（2026-05）")
inp("citi_eps28", "Citi FY2028 EPS", "NT$", None, 260.38, None, "Interested-party", UDN8, "—")
inp("ss_asic28", "賣方 FY2028 AI ASIC 營收（凱基）", "US$B", 48, 50, 52.5, "Interested-party", KGI + "：約 500 億美元（營收 67%）", "高盛 525 億美元（經濟日報 2026-08-03）、480 億美元（2026-05）")
inp("ss_hu28", "凱基假設：FY2028 v9 Humufish 出貨", "百萬顆", None, 4.0, None, "Interested-party", KGI, "—")
inp("ss_asic27", "賣方 FY2027 AI ASIC 營收（高盛，區間中點）", "US$B", 12, 14, 20.3, "Interested-party", UDN8 + "：120–160 億美元（同文另有 203 億美元）", "—")

sec("H. 價格與市場倍數")
inp("px", "股價（2026-10-22 收盤）", "NT$", None, 4010.0, None, "Interested-party", SA + "：10-22 收盤 4,010（−3.14%）", "財報狗同值；10-23 收盤：未找到")
inp("pe_ttm", "近 12 月 P/E（stockanalysis）", "x", None, 66.23, None, "Interested-party", SA, "companiesmarketcap 61.6")
inp("fpe", "前瞻 P/E（stockanalysis）", "x", None, 38.75, None, "Interested-party", SA, "—")
inp("cons_feps", "共識前瞻 12 月 EPS（現價 ÷ 前瞻 P/E）", "NT$", None, "=px/fpe", None, "Derived", "stockanalysis 前瞻 P/E 38.75", "共識前瞻 EPS 的直接數字：未找到")
inp("cons_tp", "市場共識目標價（賣方平均）", "NT$", 4000, 6067, 10369, "Interested-party", SAF + "：28 位分析師平均 6,067、中位數 5,861、區間 4,000–10,369（2026-10-23 擷取）", "stockanalysis 首頁：27 位 6,075")
inp("nim5", "近五年平均淨利率（FY2021–25；對照用）", "%", None, "=AVERAGE(111421/493415,118141/548796,76979/433446,106387/530586,105319/595966)", None, "Derived",
    "stockanalysis 財報頁：FY2021–25 營收與淨利", "十年平均（2016–2025）：未找到（2016–2020 營收與淨利未取得）；收斂法依工作單改用最後一年 × 0.8")

sec("I. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%）", "美、台、日、韓同一值")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "永續成長與 FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均，FY2023–25）", "%", None, "=AVERAGE(0.1889,0.2750,0.2606)", None, "Derived", "stockanalysis 比率頁：FY2023 18.89%、FY2024 27.50%、FY2025 26.06%", "單一來源；近 12 月 23.86%（同站）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2029 淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（FY2029 與「FY2029 × 0.8」的中點）", "%", None, 0, None, "Assumed", "十年平均未找到，依工作單以最後一年 × 0.8 代替", "—")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（FY2029 × 0.8）", "%", None, 0, None, "Assumed", "同上", "—")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    for u_, nm_ in URLNAME.items():          # 網址只列在 Sources 頁；In 頁寫來源名稱
        s1 = s1.replace(u_, nm_); s2 = s2.replace(u_, nm_)
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for k in TP_KEYS:
    for c in range(1, 10): IN.cell(INROW[k], c).fill = KEY
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點（NT$B） ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣、基本；股數為十億股）。藍字＝輸入（公司營運報告）；黑字＝公式。產品線金額＝公司公布占比 × 營收（Derived）"
Q["A2"] = "FY2026 錨點＝1Q、2Q26 實績＋3Q26 營收實績（月營收）與指引（毛利率、費用）＋4Q26 假設；AI ASIC（資料中心）全年取公司「超過 20 億美元」，1–3Q 部分全數放在 3Q 欄（簡化）。3Q26 財報約 10 月底公布，屆時以實績取代 E 欄"
hdr(Q, 3, ["項目", "FY2025A", "1Q26（實績）", "2Q26（實績）", "3Q26（營收實績＋指引）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2025：4Q25 營運報告；1Q26、2Q26：2Q26 營運報告（公司）"
def qrow(r, label, vals, fmt="#,##0.00", tag="Verified", s1=S1, s2="—", ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
# 3Q 產品線：2Q 金額 × 公司季增指引，再等比例放大到營收實績
K3 = "q3_rev/(D5*(1+q3_mob)+D6*(1+q3_sep)+D7*(1+q3_pic))"
qrow(4, "營收", {"B": 595.966, "C": 149.151, "D": 152.183, "E": "=q3_rev", "F": "=SUM(F5:F7)"}, tag="Verified／Assumed（4Q）", s2="3Q26：月營收實績（科技新報、財報狗）")
qrow(5, "  手機", {"B": "=B4*0.53", "C": "=C4*0.49", "D": "=D4*0.41", "E": f"=D5*(1+q3_mob)*{K3}", "F": "=E5*(1+q4_mob)"}, tag="Derived（占比 × 營收）／Assumed",
     s2="占比：1Q26 49%、2Q26 41%（公司）；FY2025 約 53%（公司：手機超過 100 億美元 ÷ 191 億美元；占比未揭露，Derived）")
qrow(6, "  智慧邊緣平台（含 AI ASIC）", {"B": "=B4*0.415", "C": "=C4*0.46", "D": "=D4*0.53", "E": f"=D6*(1+q3_sep)*{K3}", "F": "=(E6-E8)*(1+q4_sepx)+F8"}, tag="Derived／Assumed",
     s2="占比：1Q26 46%、2Q26 53%（公司）；FY2025 約 41.5%（Derived；公司僅揭露美元年增 21%）")
qrow(7, "  電源 IC", {"B": "=B4*0.055", "C": "=C4*0.05", "D": "=D4*0.06", "E": f"=D7*(1+q3_pic)*{K3}", "F": "=E7*(1+q4_pic)"}, tag="Derived／Assumed", s2="占比：1Q26 5%、2Q26 6%、4Q25 5%（公司）")
qrow(8, "    其中 AI ASIC（資料中心）", {"B": 0, "C": 0, "D": 0, "E": "=asic13*fx26", "F": "=(dc26-asic13)*fx26"}, tag="Interested-party／Assumed", s1="2Q26 講稿：FY2026 資料中心超過 20 億美元、4Q26 量產",
     s2="FY2025 視為 0（量產自 4Q26）；1–3Q 部分放在 3Q 欄")
qrow(9, "  智慧邊緣平台（不含 AI ASIC）", {c: f"={c}6-{c}8" for c in "BCDEF"}, tag="Derived", s1="—")
qrow(10, "毛利", {"B": 283.080, "C": 69.055, "D": 70.308, "E": "=E4*q3_gm", "F": "=q4_gmx*(F5+F9+F7)+F8*(m_core*(1-hf26)+m_hbm*hf26)"}, tag="Verified／Interested-party（3Q）／Assumed（4Q）",
     s2="4Q：非 ASIC × 毛利率＋ASIC（核心 m_core、代購 HBM m_hbm）")
qrow(11, "毛利率", {c: f"={c}10/{c}4" for c in "BCDEFG"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="公司：FY2025 47.5%、1Q26 46.3%、2Q26 46.2%")
qrow(12, "營業費用（推銷＋管理＋研發）", {"B": 179.610, "C": 46.165, "D": 47.440, "E": "=q3_opx", "F": "=q4_opx"}, tag="Verified／Interested-party／Assumed")
qrow(13, "營業利益", {c: f"={c}10-{c}12" for c in "BCDEF"}, tag="Derived", s2="公司：FY2025 103.470、1Q26 22.891、2Q26 22.868（核對見下方）")
qrow(14, "營業外淨額", {"B": 21.418, "C": 4.129, "D": 5.094, "E": "=q_nonop", "F": "=q_nonop"}, tag="Verified／Assumed")
qrow(15, "稅前利益", {c: f"={c}13+{c}14" for c in "BCDEF"}, tag="Derived", s1="—")
qrow(16, "所得稅", {"B": 18.770, "C": 2.643, "D": 3.357, "E": "=E15*q_tax", "F": "=F15*q_tax"}, tag="Verified／Assumed")
qrow(17, "非控制權益", {"B": 0.798, "C": 0.222, "D": 0.270, "E": "=q_nci", "F": "=q_nci"}, tag="Verified／Derived")
qrow(18, "歸屬母公司淨利", {c: f"={c}15-{c}16-{c}17" for c in "BCDEF"}, tag="Derived", s2="公司：FY2025 105.319、1Q26 24.154、2Q26 24.335")
qrow(19, "  公司揭露歸屬母公司淨利", {"B": 105.319, "C": 24.154, "D": 24.335}, ann="", tag="Verified")
qrow(20, "加權股數（十億股）", {"B": "=105.319/66.16", "C": "=24.154/15.17", "D": "=24.335/15.28", "E": "=q_sh", "F": "=q_sh"}, fmt="0.0000", ann="=AVERAGE(C20:F20)", tag="Derived", s1="淨利 ÷ EPS（公司）")
qrow(21, "EPS（基本）", {c: f"={c}18/{c}20" for c in "BCDEF"}, ann="=G18/G20", tag="Derived", s1="淨利 ÷ 股數")
qrow(22, "  公司揭露 EPS", {"B": 66.16, "C": 15.17, "D": 15.28}, ann="", tag="Verified")
Q.cell(24, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
checks = [("FY2026 營收年增", "=G4/B4-1", "0.0%"), ("FY2026 美元營收年增（各季營收 ÷ 匯率；對照公司高個位數）", "=(C4/fx_q1+D4/fx_q2+(E4+F4)/fx26)/co_usd25-1", "0.0%"),
          ("FY2026 手機年增", "=G5/B5-1", "0.0%"), ("FY2026 智慧邊緣平台（不含 ASIC）年增", "=G9/B9-1", "0.0%"),
          ("3Q26 營收 ÷ 指引上緣 159.8 − 1", "=E4/159.8-1", "0.0%"), ("4Q26 營收季增", "=F4/E4-1", "0.0%"),
          ("檢查：FY2025 營業利益 自算 − 公司 103.470（應約為 0）", "=B13-103.47", "0.000"), ("檢查：1H26 歸屬淨利 自算 − 公司 48.489（應約為 0）", "=C18+D18-48.489", "0.000"),
          ("檢查：FY2025 歸屬淨利 自算 − 公司（應約為 0）", "=B18-B19", "0.000")]
for i, (lab, f, fm) in enumerate(checks, 25):
    Q.cell(i, 1, lab); Q[f"B{i}"] = f; Q[f"B{i}"].number_format = fm
for nm, rr in (("rev", 4), ("mob", 5), ("sep", 6), ("pic", 7), ("asic", 8), ("sepx", 9), ("gp", 10), ("opx", 12), ("oi", 13), ("nonop", 14), ("tax", 16), ("nci", 17), ("ni", 18), ("sh", 20), ("eps", 21)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 62
for c in "BCDEFG": Q.column_dimensions[c].width = 14
Q.column_dimensions["H"].width = 30; Q.column_dimensions["I"].width = 50; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I", 2032: "J"}
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
BU = (2026, 2027, 2028, 2029); F3 = (2027, 2028, 2029)
DEFP = dict(scn="A", uf="1", shs="", shp="0", sho="0", pfc="1", fh="f_hbm", mh="m_hbm", mc="m_core", fx="fx",
            gmob27="g_mob27_{s}", gmobL="g_mobL_{s}", gsep27="g_sep27_{s}", gsepL="g_sepL_{s}", gpic27="g_pic27_{s}", gpicL="g_picL_{s}",
            gma="0", phi="phi", opxa="0", nonop="nonop_f", tax="tax", fg="fade_g{s}", fm="fade_m{s}")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    p = lambda k: P[k].format(s=s)
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 11, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    def uexp(k, y):
        if y <= 2027: return f"u_{k}_{y % 100}"
        if y == 2028: return f"u_{k}_28*um28_{s}"
        return f"u_{k}_28*(1+g29_{k})*um29_{s}"
    shx = {"tpu": P["shp"], "oth": P["sho"]}
    sec_("A. AI ASIC（由下而上；D 欄＝FY2026 校準年；FY(t)＝CY(t)）")
    for k in XP:
        row("u_" + k, f"{PNAME[k]} 市場顆數（含情境係數）", "百萬顆", {FC[y]: "=" + uexp(k, y) for y in BU}, "0.000")
    for k in XP:
        row("s_" + k, f"聯發科份額：{PNAME[k]}", "%", {FC[y]: (f"=mtk_{k}_26" if y == 2026 else f"=MIN(1,MAX(0,mtk_{k}_{min(y, 2028) % 100}{P['shs'] if k == 'tpu' else ''}+{shx[k]}))") for y in BU}, "0.0%")
    for k in XP:
        row("b_" + k, f"聯發科顆數：{PNAME[k]}", "百萬顆", {FC[y]: f"={FC[y]}{R['u_' + k]}*{FC[y]}{R['s_' + k]}" + ("" if y == 2026 else f"*{P['uf']}") for y in BU}, "0.000")
    for k in XP:
        row("g_" + k, f"每顆 HBM：{PNAME[k]}", "GB", {FC[y]: f"=gb_{k}_{min(y, 2028) % 100}" for y in BU}, "0")
    row("hpx", "HBM 單價", "US$/GB", {FC[y]: f"=hbm_{s}_{y % 100}" for y in BU})
    for k in XP:
        row("c_" + k, f"每顆核心售價：{PNAME[k]}", "US$K", {FC[y]: f"=core_{k}" + (f"*{P['pfc']}" if y >= 2028 else "") for y in BU})
    for k in XP:
        row("h_" + k, f"每顆 HBM 金額（全部）：{PNAME[k]}", "US$K", {FC[y]: f"={FC[y]}{R['g_' + k]}*{FC[y]}{R['hpx']}/1000" for y in BU})
    for k in XP:
        row("a_" + k, f"每顆聯發科營收（核心＋代購 HBM）：{PNAME[k]}", "US$K", {FC[y]: f"={FC[y]}{R['c_' + k]}+{P['fh']}*{FC[y]}{R['h_' + k]}" for y in BU},
            note="核心售價 ＋ 代購比例 × HBM 金額（未乘 k_x）" if k == "tpu" else "")
    row("xraw", "AI ASIC 營收（未校準）", "US$B", {FC[y]: "=" + "+".join(f"{FC[y]}{R['b_' + k]}*{FC[y]}{R['a_' + k]}" for k in XP) for y in BU})
    row("kx", "ASIC 校準係數 k_x ＝ FY2026 公司資料中心營收 ÷ 未校準", "x", {"D": f"=dc26/D{r - 1}"}, "0.000",
        note="Derived：以公司 FY2026「超過 20 億美元」校準；FY2027 起沿用。顆數、份額、單價的合併係數")
    row("asu", "AI ASIC 營收（美元）", "US$B", {FC[y]: f"={FC[y]}{R['xraw']}*$D${R['kx']}" for y in BU})
    row("fxr", "匯率", "NT$/US$", {"D": "=fx26", **{FC[y]: f"={P['fx']}" for y in F3}})
    row("asic", "AI ASIC 營收", "NT$B", {"C": "=a25_asic", **{FC[y]: f"={FC[y]}{R['asu']}*{FC[y]}{R['fxr']}" for y in BU}})
    row("xh", "  其中代購 HBM", "NT$B", {FC[y]: f"=$D${R['kx']}*{P['fh']}*(" + "+".join(f"{FC[y]}{R['b_' + k]}*{FC[y]}{R['h_' + k]}" for k in XP) + f")*{FC[y]}{R['fxr']}" for y in BU})
    row("xc", "  其中核心（設計＋晶片＋封裝）", "NT$B", {FC[y]: f"={FC[y]}{R['asic']}-{FC[y]}{R['xh']}" for y in BU})
    row("xu", "聯發科 ASIC 顆數合計", "百萬顆", {FC[y]: "=" + "+".join(f"{FC[y]}{R['b_' + k]}" for k in XP) for y in BU}, "0.000")
    row("ppu", "每顆聯發科 ASIC 營收（校準後）", "US$K", {FC[y]: f"=IF({FC[y]}{R['xu']}>0,{FC[y]}{R['asu']}/{FC[y]}{R['xu']},0)" for y in BU})
    sec_("B. 手機、智慧邊緣平台（不含 ASIC）、電源 IC（FY2027 起簡化成長率）")
    for key, lab, g27, gL in (("mob", "手機", "gmob27", "gmobL"), ("sepx", "智慧邊緣平台（不含 ASIC）", "gsep27", "gsepL"), ("pic", "電源 IC", "gpic27", "gpicL")):
        row(key, lab, "NT$B", {"C": f"=a25_{key}", "D": f"=a26_{key}", "E": f"=D{r}*(1+{p(g27)})", "F": f"=E{r}*(1+{p(gL)})", "G": f"=F{r}*(1+{p(gL)})"})
    row("nas", "非 ASIC 合計", "NT$B", {FC[y]: f"={FC[y]}{R['mob']}+{FC[y]}{R['sepx']}+{FC[y]}{R['pic']}" for y in FYS})
    row("rev", "營收", "NT$B", {FC[y]: f"={FC[y]}{R['nas']}+{FC[y]}{R['asic']}" for y in FYS})
    row("g", "營收年增", "%", {FC[y]: f"={FC[y]}{R['rev']}/{FC[y - 1]}{R['rev']}-1" for y in FYS[1:]}, "0.0%")
    row("aish", "AI ASIC ÷ 營收", "%", {FC[y]: f"={FC[y]}{R['asic']}/{FC[y]}{R['rev']}" for y in FYS}, "0.0%")
    row("mobsh", "手機 ÷ 營收", "%", {FC[y]: f"={FC[y]}{R['mob']}/{FC[y]}{R['rev']}" for y in FYS}, "0.0%")
    sec_("C. 毛利與損益")
    row("gpa", "AI ASIC 毛利 ＝ m_core × 核心 ＋ m_hbm × 代購 HBM", "NT$B", {FC[y]: f"={P['mc']}*{FC[y]}{R['xc']}+{P['mh']}*{FC[y]}{R['xh']}" for y in BU})
    row("gmx", "非 ASIC 毛利率（FY2026 殘差，之後沿用＋調整）", "%", {"C": f"=a25_gp/C{R['nas']}", "D": f"=(a26_gp-D{R['gpa']})/D{R['nas']}", **{FC[y]: f"=$D{r}+{P['gma']}" for y in F3}}, "0.00%",
        note="Derived：讓 FY2026 毛利等於錨點")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{FC[y]: f"={FC[y]}{R['gmx']}*{FC[y]}{R['nas']}+{FC[y]}{R['gpa']}" for y in F3}})
    row("gm", "毛利率", "%", {FC[y]: f"={FC[y]}{R['gp']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("opx", "營業費用", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{FC[y]: f"={FC[y - 1]}{r}*(1+MAX(g_opx_min,{P['phi']}*{FC[y]}{R['g']}))+{P['opxa']}*{FC[y]}{R['rev']}" for y in F3}})
    row("oi", "營業利益", "NT$B", {FC[y]: f"={FC[y]}{R['gp']}-{FC[y]}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {FC[y]: f"={FC[y]}{R['oi']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("nonop", "營業外淨額", "NT$B", {"C": "=a25_nonop", "D": "=a26_nonop", **{FC[y]: f"={P['nonop']}" for y in F3}})
    row("pti", "稅前利益", "NT$B", {FC[y]: f"={FC[y]}{R['oi']}+{FC[y]}{R['nonop']}" for y in FYS})
    row("tax", "所得稅", "NT$B", {"C": "=a25_tax", "D": "=a26_tax", **{FC[y]: f"={FC[y]}{R['pti']}*{P['tax']}" for y in F3}})
    row("nci", "非控制權益", "NT$B", {"C": "=a25_nci", "D": "=a26_nci", **{FC[y]: "=nci" for y in F3}})
    row("ni", "歸屬母公司淨利", "NT$B", {FC[y]: f"={FC[y]}{R['pti']}-{FC[y]}{R['tax']}-{FC[y]}{R['nci']}" for y in FYS})
    row("shr", "加權股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{FC[y]: f"={FC[y - 1]}{r}*(1+sh_g)" for y in F3}}, "0.0000")
    row("eps", "EPS（基本）", "NT$", {FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['shr']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {FC[y]: f"={FC[y]}{R['eps']}/{FC[y - 1]}{R['eps']}-1" for y in FYS[1:]}, "0.0%")
    row("nim", "淨利率（歸屬母公司淨利 ÷ 營收）", "%", {FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    for y in FADE:      # FY2030–32：收斂法（上游只到 2029；非驅動因子推導）
        c, pc, k = FC[y], FC[y - 1], y - 2029
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+({p('fg')}-$G{R['g']})*{k}/3"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({p('fm')}-$G{R['nim']})*{k}/3"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['shr']}"] = f"={pc}{R['shr']}*(1+sh_g)"
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['shr']}"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("g", "0.0%"), ("rev", "#,##0.00"), ("nim", "0.00%"), ("ni", "#,##0.00"), ("shr", "0.0000"), ("eps", "#,##0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fm; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 11, "FY2030 起為收斂法：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 聯發科損益（十億新台幣；EPS 新台幣、基本；歸屬母公司）。FY2025A 實績、FY2026 錨點（三情境相同），FY2027–FY2029 由驅動因子推導，FY2030–FY2032 為收斂法。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數基準；手機 FY2027 +5%、之後 +4%；智慧邊緣 +8%／+7%；電源 IC +8%",
         "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85；手機 0%／+2%；智慧邊緣 +4%；電源 IC +4%",
         "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6；手機 −8%／−3%；智慧邊緣 −3%／+1%；電源 IC 0%"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("asic", "  AI ASIC", "NT$B", "#,##0.0"), ("mob", "  手機", "NT$B", "#,##0.0"),
        ("sepx", "  智慧邊緣平台（不含 ASIC）", "NT$B", "#,##0.0"), ("pic", "  電源 IC", "NT$B", "#,##0.0"),
        ("gm", "毛利率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"), ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("nim", "  淨利率", "%", "0.00%"),
        ("eps", "EPS", "NT$", "#,##0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{SUMR[('A', 'nim')]}")
IN.cell(INROW["fade_mB"], 5, f"=PL!G{SUMR[('B', 'nim')]}*0.9")
IN.cell(INROW["fade_mC"], 5, f"=PL!G{SUMR[('C', 'nim')]}*0.8")
for k in ("fade_mA", "fade_mB", "fade_mC"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 錨點").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（NT$B）", "+".join(f"ABS(D{PB[s]['rev']}-a26_rev)" for s in SC)),
               ("  三情境：EPS 差（NT$）", "+".join(f"ABS(D{PB[s]['eps']}-a26_eps)" for s in SC)),
               ("  情境 A：ASIC 由下而上 − Q 頁 ASIC（NT$B）", f"D{PB['A']['asic']}-a26_asic")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 10; PL.column_dimensions["K"].width = 70
for c in "CDEFGHIJ": PL.column_dimensions[c].width = 13
PL.cell(2, 1, "灰底＝FY2030–FY2032 收斂法（上游顆數只到 CY2029）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 I 節），非驅動因子推導")

# ---------------- Units：聯發科 ASIC 每顆規格、份額加總 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 聯發科 AI ASIC 規格（情境 A）。ASIC 核心售價、HBM 代購比例取自 Broadcom 模型（本專案權威，唯讀）；顆數：百萬顆（日曆年；與 HBM、台積電、Broadcom、Marvell 模型同一份）；每顆價：US$K"
A_ = PB["A"]; cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
UN.cell(ur, 1, "1. Google TPU（情境 A；聯發科營收口徑＝核心售價＋代購 HBM，乘 k_x）").font = BOLD; ur += 1
for key, lab, unit, f, fmt, note in (
        ("u", "Google TPU 市場顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['u_tpu']}", "0.00", "tsmc 成品 In 頁"),
        ("s", "聯發科份額：Google TPU", "%", lambda y: f"=PL!{FC[y]}{A_['s_tpu']}", "0.0%", "In 頁 B 節（Assumed）"),
        ("b", "聯發科顆數：Google TPU", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['b_tpu']}", "0.000", "市場層顆數 × 份額"),
        ("h", "每顆 HBM 金額：Google TPU", "US$K", lambda y: f"=PL!{FC[y]}{A_['h_tpu']}", "0.00", "容量 × HBM 單價（情境 A）"),
        ("p", "每顆聯發科營收（校準後）：Google TPU", "US$K", lambda y: f"=PL!{FC[y]}{A_['a_tpu']}*PL!$D${A_['kx']}", "0.00", "（核心＋代購比例 × HBM）× k_x"),
        ("asu", "聯發科 AI ASIC 營收（美元）", "US$B", lambda y: f"=PL!{FC[y]}{A_['asu']}", "0.00", "含其他 ASIC（基準份額 0）")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit); UR[key] = ur
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); ur += 1
ur += 1
UN.cell(ur, 1, "2. 份額加總檢查（Broadcom＋Marvell＋聯發科 ≤ 100%；Broadcom、Marvell 份額讀各自成品 In 頁）").font = BOLD; ur += 1
for k in XP:
    UN.cell(ur, 1, f"Broadcom＋Marvell＋聯發科：{PNAME[k]}"); UN.cell(ur, 2, "%")
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f"=av_{k}_{min(y, 2028) % 100}+mv_{k}_{min(y, 2028) % 100}+PL!{FC[y]}{A_['s_' + k]}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
    UN.cell(ur, 7, "其餘為世芯、創意、華為等"); UR["sum_" + k] = ur; ur += 1
UN.cell(ur, 1, "聯發科 TPU 份額上限（1 − Broadcom − Marvell 基準）"); UN.cell(ur, 2, "%")
for y in BU:
    UN[f"{cyc(y)}{ur}"] = f"=1-av_tpu_{min(y, 2028) % 100}-mv_tpu_{min(y, 2028) % 100}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
UN.cell(ur, 7, "In 頁份額高值受此上限限制"); UR["cap"] = ur; ur += 2
UN.cell(ur, 1, "3. 校準值（只用 FY2026 已發生的實績、當季指引與公司全年資料中心說法；不用共識、不用 FY2027 起的目標）").font = BOLD; ur += 1
for key, lab, f, fmt, note in (("kx", "ASIC 校準係數 k_x", f"=PL!D{A_['kx']}", "0.000", "公司 FY2026 資料中心營收 ÷ 未校準由下而上"),
                               ("gmx", "非 ASIC 毛利率（FY2026 殘差）", f"=PL!D{A_['gmx']}", "0.00%", "讓 FY2026 毛利等於錨點"),
                               ("dc26", "FY2026 資料中心營收（公司）", "=dc26", "0.00", "US$B")):
    UN.cell(ur, 1, lab); UN[f"C{ur}"] = f; UN[f"C{ur}"].number_format = fmt; UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：聯發科 ASIC 營收受市場層 TPU 顆數與 Broadcom、Marvell 份額共同限制；CY2028 份額上限 20%（1 − 75% − 5%）。賣方 FY2028 約 500 億美元的 ASIC 營收需要的顆數或單價遠高於此口徑（Recon 第 B 節）。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 56; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 13
UN.column_dimensions["G"].width = 56

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」；目標價＝DCF 與 Comparable 加權（第 0 節）；十年 P/E 規則值（第 3 節，年底落後 P/E）只作參考"
V["A2"] = "循環位置：聯發科不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 Amkor、廣達；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 51
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-22 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "#,##0"
for k, y, lab in ((3, 2026, "現價 ÷ FY2026 EPS（錨點）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 B）≈ 0.2 × FY2026＋0.8 × FY2027"); V[f"B{O + 7}"] = f"=0.2*PL!D{SUMR[('B', 'eps')]}+0.8*PL!E{SUMR[('B', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 B）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（stockanalysis）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 3}"] = "FY2026 三情境相同（錨點）"; V[f"E{O + 4}"] = "FY2027 三情境只差非 ASIC 成長率（顆數情境係數從 CY2028 開始）"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, CY0 + 1, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
# 十年 P/E：companiesmarketcap 年底 P/E；自算＝digrin 12 月未調整股價 ÷ 公司 EPS（2016 EPS 未找到）
PE_CMC = [9.22, 12.6, 11.5, 21.6, 20.7, 12.7, 6.93, 19.2, 20.4, 21.5]
PX_DEC = [216.5, 294.0, 229.5, 443.5, 747.0, 1190.0, 625.0, 1015.0, 1415.0, 1430.0]
EPS_H = [None, 15.56, 13.26, 14.69, 26.01, 70.56, 74.59, 48.51, 66.92, 66.16]
PE_SELF = [f"={p_}/{e_}" if e_ else None for p_, e_ in zip(PX_DEC, EPS_H)]
T10 = vr + 2
V.cell(T10, 1, "5. 十年 P/E（2016–2025 年底；兩來源平均，分位排名用；無負值或 >200 需剔除）").font = BOLD
hdr(V, T10 + 1, ["年底", "兩來源平均 P/E", "companiesmarketcap", "自算"])
RB0 = T10 + 16
for i, y in enumerate(range(2016, 2026)):
    rr = T10 + 2 + i
    V.cell(rr, 1, str(y))
    V.cell(rr, 2, f"=AVERAGE(C{rr}:D{rr})").number_format = "0.00"
    V.cell(rr, 3, f"=B{RB0 + 5 + i}").number_format = "0.00"
    V.cell(rr, 4, f'=IF(C{RB0 + 5 + i}="","",C{RB0 + 5 + i})').number_format = "0.00"
PE10 = f"$B${T10 + 2}:$B${T10 + 11}"
V.cell(T10 + 12, 1, "年數"); V[f"B{T10 + 12}"] = f"=COUNT({PE10})"
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2016–2025 共 10 年；早年價格疑為除息調整後，倍數偏低）"},
    {"label": "自算（digrin 12 月未調整股價 ÷ 公司 EPS）", "vals": PE_SELF, "tag": "Derived（價格為月資料的 12 月值；EPS：2017 科技新報、2018–2025 HiStock 與公司；2016 EPS 未找到）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。stockanalysis 2021–2025 年底 P/E（16.99、8.42、20.99、21.20、21.67）與自算相同，只作核對。聯發科為高配息公司，companiesmarketcap 早年倍數偏低（疑為除息調整價格）；十年皆為手機 SoC 期，未含 AI ASIC 放量後的成長溢價，套用在 FY2027 EPS 時偏向低估。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025", (V.cell(RB0 + 5, 1).value, V.cell(RB0 + 14, 1).value)
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# 0. 目標價（Andy 2026-10-25 11:28／11:53：DCF 與 Comparable 各半；列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期）
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "#,##0.00"); EPSR[y] = 8 + i
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)" for c in "BCD"], "#,##0.00")
def dcf_formula(c, r="tp_r", g="tp_g", reinv="$B$15"):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], "#,##0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "#,##0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "#,##0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("高通（Qualcomm）", "QCOM", 18.98, "2026-10-23"), ("博通（Broadcom）", "AVGO", 18.52, "2026-10-23"), ("邁威爾（Marvell）", "MRVL", 40.44, "2026-10-23"),
         ("世芯-KY（Alchip）", "3661.TW", 17.15, "2026-10-22"), ("創意（GUC）", "3443.TW", 102.85, "2026-10-08")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "手機 SoC（高通）＋ASIC 設計（博通、邁威爾、世芯、創意）；不含 A 股（非 A 股公司，Andy 2026-10-25 11:53）")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(2/12*B{EPSR[2027]}+10/12*B{EPSR[2028]})+tp_wB*(2/12*C{EPSR[2027]}+10/12*C{EPSR[2028]})+tp_wC*(2/12*D{EPSR[2027]}+10/12*D{EPSR[2028]})", None, None], "#,##0.00",
     "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "#,##0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "", "", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "#,##0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "#,##0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "#,##0.00", "stockanalysis：28 位分析師平均 6,067（中位數 5,861），2026-10-23 擷取")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-22 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝現價 ÷ 前瞻 P/E（In 頁 cons_feps）")
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
tset(46, "band", "價格帶（DCF 不如預期～順利）", ['=TEXT(D23,"#,##0")&"～"&TEXT(B23,"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 5 節兩來源平均）的分位；高於十年最大值記為 100%；只作檢查，不進目標價")
V.cell(48, 1, "目標價方法")
V["B48"] = ("DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為近三年平均約 24%（再投資率約 17%），FY2027–32 明確期＋永續；聯發科為輕資產 IC 設計，不改用淨利＋折舊－資本支出。"
            "FY2030 起為收斂法，ASIC 第二、第三顆的成長未在明確期展開，DCF 偏低估的方向。"
            "Comparable：高通、博通、邁威爾、世芯、創意前瞻 P/E 中位數（高通 19x、博通 18.5x 落在中間；創意 103x 為極端值，中位數不受影響）。"
            "主要偏誤：ASIC 營收受市場層 TPU 顆數與 Broadcom 份額限制，遠低於賣方（Recon 第 B 節），本模型 EPS 與目標價都偏向低估賣方路徑。"
            "情境對應：A＝順利（顆數基準）、B＝平平（顆數 ×0.8／0.85）、C＝不如預期（顆數 ×0.6、手機下滑）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；FY2026 恆等於錨點，k_x 與非 ASIC 毛利率在每個案例內重算）：對 FY2027、FY2028 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("TPU 份額 −5pt", "CY2027 起聯發科占 TPU −5 個百分點", dict(shp="(-0.05)")), ("TPU 份額取高值（上限）", "CY2027 20%、CY2028 起 20%", dict(shs="_hi")),
         ("ASIC 顆數 ×0.8", "CY2027 起聯發科顆數 ×0.8", dict(uf="uf_lo")), ("ASIC 顆數 ×1.2", "×1.2（超出份額上限，僅示意）", dict(uf="uf_hi")),
         ("第二顆 ASIC 單價 ×2", "CY2028 起核心售價 ×2（賣方：新世代單價數倍）", dict(pfc="pfc_hi")), ("核心售價 ×0.8", "CY2028 起 ×0.8", dict(pfc="pfc_lo")),
         ("其他 ASIC 份額 +5pt", "CY2027 起第二客戶（基準 0）", dict(sho="0.05")),
         ("HBM 代購比例 低", "代購比例取低值（k_x 重算）", dict(fh="f_hbm_lo")), ("HBM 代購比例 高", "代購比例取高值", dict(fh="f_hbm_hi")),
         ("ASIC 核心毛利率 35%", "基準 45%", dict(mc="m_core_lo")), ("ASIC 核心毛利率 55%", "", dict(mc="m_core_hi")),
         ("手機 情境 C 路徑", "FY2027 −8%、之後 −3%", dict(gmob27="g_mob27_C", gmobL="g_mobL_C")),
         ("智慧邊緣 情境 C 路徑", "FY2027 −3%、之後 +1%", dict(gsep27="g_sep27_C", gsepL="g_sepL_C")),
         ("非 ASIC 毛利率 −2pt", "FY2027 起", dict(gma="gmadj_lo")), ("營業費用彈性 0.6", "基準 0.3", dict(phi="phi_hi")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi")), ("匯率 30", "基準 32（ASIC 美元營收換算）", dict(fx="fx_lo")),
         ("稅率 17%", "基準 14%", dict(tax="tax_hi")), ("營業外 10", "基準 18", dict(nonop="nonop_f_lo"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0"), ("D", "gm", 2027, "0.00%"), ("E", "eps", 2027, "#,##0.00"),
                           ("G", "rev", 2028, "#,##0"), ("H", "gm", 2028, "0.00%"), ("I", "eps", 2028, "#,##0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"F{i}"] = f"=E{i}/E$4-1"; SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c in "FJ": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 B", "CY2028 顆數 ×0.8、CY2029 ×0.85；非 ASIC 成長較低", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "CY2028–29 顆數 ×0.6；手機 −8%", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 三情境只差非 ASIC 成長率；顆數情境 B、C 的差異從 FY2028 起。TPU 份額高值受 Broadcom、Marvell 份額上限限制；第二顆 ASIC 單價 ×2 是賣方路徑的主要差異來源。")
SENS_NOTE = nn + 3
TPS0 = nn + 5
SE.cell(TPS0, 1, "目標價參數（只改 DCF 參數；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "#,##0.00"
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.3", "", dict(w="0.3")), ("DCF 權重 0.7", "", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
assert TPS_LAST < 40, TPS_LAST
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 44
for c in "CDEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引與目標、共識、券商。只作事後對照，不為對上數字而調參數。公司 FY2027 SAM 份額目標與券商 FY2027–28 估計為 Interested-party，不用於校準"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. FY2026 錨點 vs 公司展望與共識").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "stockanalysis（25 位）", "凱基 678.0；3Q26 營收已超過指引上緣 5%")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "stockanalysis（25 位）", "共識區間 58.48–73.35", fmt="#,##0.00")
rec("FY2026 美元營收年增 vs 公司高個位數（百分點）", "=co_g26", lambda s: "=Q!B26", "2Q26 講稿", "3Q26 營收超出指引、4Q26 ASIC 量產", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "B. AI ASIC vs 公司目標與賣方（市場層顆數 × 份額口徑）").font = BOLD; rr_ += 1
rec("FY2027 ASIC（美元）vs 公司 SAM × 份額目標中點（80 × 17.5%）", "=co_sam27*co_sh27", pl("asu", 2027), "2Q26 講稿", "目標區間 12–16（US$B）；份額目標的分母為 SAM，不是 TPU 顆數", fmt="0.00")
rec("FY2027 ASIC（美元）vs 高盛 120–160 億美元中點", "=ss_asic27", pl("asu", 2027), "經濟日報 2026-08-03", "", fmt="0.00")
rec("FY2027 ASIC ÷ 營收 vs 公司「20% 可能」（百分點）", "=co_dc27sh", pl("aish", 2027), "4Q25 法說", "", fmt="0.0%", pt=True)
rec("FY2028 ASIC（美元）vs 凱基 500 億美元", "=ss_asic28", pl("asu", 2028), "凱基 2026-07-23", "差距為本模型最大的口徑問題（見下兩列）", fmt="0.00")
rec("FY2028：達到凱基 ASIC 所需聯發科顆數（＝500 億 ÷ 模型每顆營收）", None, lambda s: f"=ss_asic28/PL!F{PB[s]['ppu']}", "推導", "百萬顆；凱基假設 v9 Humufish 400 萬顆", fmt="0.00")
rec("FY2028：所需顆數 ÷ 市場層 TPU 顆數", None, lambda s: f"=ss_asic28/PL!F{PB[s]['ppu']}/PL!F{PB[s]['u_tpu']}", "推導", ">（1 − Broadcom 75% − Marvell 5%）＝ 20% 即與 avgo 份額口徑衝突", fmt="0.0%")
rec("FY2028：凱基顆數下所需每顆營收（＝500 億 ÷ 400 萬顆）÷ 模型每顆營收", None, lambda s: f"=ss_asic28/ss_hu28/PL!F{PB[s]['ppu']}", "推導", "倍數；凱基稱 v9 單價至少 +200%", fmt="0.00")
RC.cell(rr_, 1, "C. EPS vs 券商（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 EPS vs 凱基", "=kgi_eps27", pl("eps", 2027), "凱基 2026-07-23", "", fmt="#,##0.00")
rec("FY2027 EPS vs Citi", "=citi_eps27", pl("eps", 2027), "經濟日報 2026-08-03", "", fmt="#,##0.00")
rec("FY2028 EPS vs 凱基", "=kgi_eps28", pl("eps", 2028), "凱基 2026-07-23", "差距 >30%：見 README", fmt="#,##0.00")
rec("FY2028 EPS vs Citi", "=citi_eps28", pl("eps", 2028), "經濟日報 2026-08-03", "", fmt="#,##0.00")
rec("FY2027 營收 vs 凱基", "=kgi_rev27", pl("rev", 2027), "凱基 2026-07-23", "", fmt="#,##0.0")
rec("共識前瞻 12 月 EPS vs 模型（情境加權 2／12 × FY2027＋10／12 × FY2028）", "=cons_feps", lambda s: f"=Val!B{T['feps']}", "stockanalysis 前瞻 P/E 推得", "三欄同值（情境加權）", fmt="#,##0.00")
RC.cell(rr_, 1, "D. 校準與份額").font = BOLD; rr_ += 1
rec("k_x（ASIC 校準係數）", None, lambda s: f"=PL!D{PB[s]['kx']}", "公司 FY2026 資料中心 ÷ 未校準", "接近 1＝市場層 TPU 顆數 × 5% × Broadcom 模型單價與公司說法一致", fmt="0.000")
rec("非 ASIC 毛利率（FY2026 殘差）", None, lambda s: f"=PL!D{PB[s]['gmx']}", "FY2026 毛利殘差", "FY2025 47.5%（全部非 ASIC）", fmt="0.00%")
rec("CY2028 Broadcom＋Marvell＋聯發科 TPU 份額", 1.0, lambda s: f"=Units!E{UR['sum_tpu']}", "avgo、mrvl 成品", "應 ≤ 100%", fmt="0.0%", pt=True)
rec("FY2027 營業利益率 vs 高盛 22%（百分點）", 0.22, pl("oim", 2027), "經濟日報 2026-05-03", "", fmt="0.0%", pt=True)
rec("FY2028 營業利益率 vs 高盛 33%（百分點）", 0.33, pl("oim", 2028), "經濟日報 2026-05-03", "", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "E. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以聯發科 ASIC 營收 ÷ 部署 GW 對照 Tokenomics 非 NVIDIA 加速器比例"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 80
for c in "BCDEFGH": RC.column_dimensions[c].width = 11
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60
RECON_LAST = rr_ - 1

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝聯發科法說資料；券商、資料商、媒體皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：2Q26 營運報告（1Q26、2Q26 損益表、資產負債）", R2Q),
        ("公司：2Q26 法說講稿（3Q26 指引、資料中心超過 20 億美元、SAM 800 億美元、份額 15–20%、第二顆 ASIC 2028）", P2Q),
        ("公司：1Q26 法說逐字稿（ASIC 數十億美元、SAM 700–800 億美元、份額 10–15%）", T1Q),
        ("公司：4Q25 營運報告（FY2025、4Q25、3Q25 損益表）", R4Q),
        ("公司：4Q25 法說逐字稿（FY2025 美元營收 191 億、資料中心 2026 超過 10 億美元、2027 占營收 20%）", T4Q),
        ("媒體：科技新報（9 月營收、前 9 月累計）", SEP26), ("媒體：財報狗（3Q26 營收 1,680.58 億、4Q 展望）", SDOG), ("媒體：風傳媒（2Q26 法說摘要）", STORM),
        ("券商：凱基 2026-07-23（FY2026–28 預估、v8x／v9／v10 TPU、FY2028 ASIC 500 億美元、目標價 6,150）", KGI),
        ("媒體：經濟日報 2026-08-03（高盛、Citi 目標價與 EPS、ASIC 2026–28）", UDN8), ("媒體：經濟日報 2026-05-03（高盛 EPS、ASIC 單價數倍、25x P/E）", UDN5),
        ("資料商：stockanalysis（現價 10-22、近 12 月與前瞻 P/E、ROE、FY2021–25 財務）", SA), ("資料商：stockanalysis 共識頁（FY2026 共識、目標價）", SAF),
        ("資料商：stockanalysis 同業前瞻 P/E（QCOM、AVGO、MRVL、3661、3443）", "https://stockanalysis.com/stocks/qcom/；https://stockanalysis.com/stocks/avgo/；https://stockanalysis.com/stocks/mrvl/；https://stockanalysis.com/quote/tpe/3661/；https://stockanalysis.com/quote/tpe/3443/"),
        ("資料商：companiesmarketcap 年底 P/E", "https://companiesmarketcap.com/mediatek/pe-ratio/"),
        ("資料商：digrin 月股價（12 月未調整值）", "https://www.digrin.com/stocks/detail/2454.TW/price"),
        ("資料商：HiStock 歷年 EPS（2018–2025）；科技新報 2017 EPS 15.56", "https://histock.tw/stock/2454/%E6%AF%8F%E8%82%A1%E7%9B%88%E9%A4%98；https://finance.technews.tw/2018/01/31/mediatek-2017q4/"),
        ("資料商：財報狗（現價 10-22 核對）", "https://statementdog.com/analysis/2454"),
        ("市場層：tsmc、avgo、mrvl 成品（本 repo，唯讀）", f"{rel(TSMC_F)}；{rel(AVGO_F)}；{rel(MRVL_F)}"),
        ("未能讀取（權限請求逾時）：公司 2Q26 新聞稿、4Q25 講稿 PDF、finanzen.ch 歷年 EPS", "https://www.mediatek.com/hubfs/MediaTek%20Assets/Pdfs/Quarterly%20Earnings%20Release/2026/Quarterly%20Earnings%20Release-2026Q2/Press%20Release.pdf"),
        ("待補（人工）：3Q26 財報（約 10 月底）；客戶名稱與 TPU 世代分工；聯發科 ASIC 每顆單價與 HBM 採購方式；2016 年 EPS 與 2016–2020 淨利率（十年平均）；前瞻 P/E 十年歷史（Bloomberg）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("聯發科（2454.TW）收支與估值模型 v0.1.1（AI 半導體 Project：運算層，手機 SoC＋AI ASIC）　2026-10-25", True),
 ("版本：v0.1（初版）。ASIC 每顆規格（核心售價、HBM 代購）讀 Broadcom 模型（本專案權威），份額與 Broadcom、Marvell 合計不超過 100%；手機、智慧邊緣平台、電源 IC 用簡化成長率。本版未做先驗。", False),
 ("模型命題：聯發科的 AI ASIC 營收＝Google TPU 顆數（市場層）× 聯發科份額 × 每顆營收（Broadcom 模型規格）；FY2026 只有約 20 億美元，賣方卻估 FY2028 達 500 億美元。在 Broadcom 占 TPU 75%、Marvell 5% 的份額口徑下，聯發科最多拿到兩成 TPU，推得出多少 ASIC 營收與 EPS？", False),
 ("核心問題：以市場層 TPU 顆數與份額上限往前推，聯發科 FY2027–28 的 ASIC 營收與 EPS 離公司目標、賣方估計有多遠；差距要靠顆數、份額還是單價補上，現價約 39 倍前瞻 P/E 隱含哪一條路徑。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。FY2025A 實績、FY2026 錨點、FY2027–FY2029 驅動因子推導、FY2030–FY2032 收斂法。金額：十億新台幣；EPS：新台幣（基本）。口徑：歸屬母公司淨利（TIFRS）。3Q26 財報約 10 月底公布，本版只納入 3Q26 營收實績。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI ASIC：tsmc 成品的 Google TPU 顆數 × 情境係數（CY2028–29）× 聯發科份額（Assumed；CY2026 5%、CY2027 15%、CY2028 起 18%，高值以 1 − Broadcom − Marvell 份額為上限）×（核心售價 8K × 第二顆 ASIC 單價係數＋代購比例 0.6 × HBM 金額）× k_x（FY2026 以公司「超過 20 億美元」校準）× 匯率 32。", False),
 ("非 ASIC：手機、智慧邊緣平台（不含 ASIC）、電源 IC 由 FY2026 錨點依情境成長率推（A 順利／B 平平／C 不如預期）。", False),
 ("毛利：ASIC 核心毛利率 45%（Assumed）、代購 HBM 10%（Broadcom 模型）；非 ASIC 毛利率為 FY2026 殘差（Derived）後沿用。營業費用彈性 0.3（ASIC 營業費用低）、年增下限 6%；營業外 18、稅率 14%、股數持平。", False),
 ("校準：FY2026 ＝ 1H26 實績＋3Q26 營收實績（月營收）＋3Q26 指引（毛利率 46%、費用 48.5）＋4Q26 假設（手機 −5%、智慧邊緣 −5%、ASIC 約 17 億美元量產）。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2）＋0.5 × Comparable 目標價（高通、博通、邁威爾、世芯、創意前瞻 P/E 中位數 × 12 個月後前瞻 EPS；不含 A 股）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解。十年 P/E 規則值 × FY2027 EPS 保留為參考（不是目標價）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。十年平均淨利率未找到，B、C 依工作單以 FY2029 × 0.8 代替。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司數字取自聯發科法說營運報告與講稿；券商與資料商為 Interested-party。", False),
 ("", False),
 ("已知限制：(1) 公司未揭露 ASIC 客戶、每顆單價、HBM 採購方式，TPU 份額與單價為 Assumed；(2) ASIC 份額受 avgo 的 Broadcom TPU 份額（CY2028 75%）限制，賣方路徑（FY2028 約 500 億美元）與本專案市場層口徑不相容，需由 avgo 或市場層修訂，不在本模型自改；(3) 第二顆 ASIC 單價數倍（賣方說法）只放 Sens；(4) FY2025 產品線占比為 Derived；(5) 3Q26 毛利與費用為指引值；(6) 十年 P/E 為手機 SoC 期，未含 ASIC 成長溢價；2016 EPS 未找到；(7) R4 產業部署路徑未建（Recon 預留列）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170


# ---------------- v0.1.1：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00、14:17）；各頁只在現有最後一列之後追加，不位移既有列 ----------------
_up = next(r for r in range(1, V.max_row + 1) if V.cell(r, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 2; REC_ROW = GRADE_ROW + 1
V.cell(GRADE_ROW, 1, "上檔空間級距")
V.cell(GRADE_ROW, 2, f'=IF(B{_up}>0.5,"大於 50%",IF(B{_up}>0.25,"25%–50%",IF(B{_up}>=0,"0%–25%",IF(B{_up}>=-0.25,"0%～−25%",IF(B{_up}>=-0.5,"−25%～−50%","小於 −50%")))))')
V.cell(GRADE_ROW, 5, f"以「上檔空間：加權目標價」（B{_up}）分六級（Andy 2026-10-25 14:00）")
V.cell(REC_ROW, 1, "投資建議")
V.cell(REC_ROW, 2, f'=IF(B{_up}>0.5,"強力買進",IF(B{_up}>0.25,"買進",IF(B{_up}>=-0.25,"中立",IF(B{_up}>=-0.5,"放空","強力放空"))))')
V.cell(REC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = ['目標價 NT$2,057（DCF 1,846、Comparable 2,268 各半），比現價 4,010（2026-10-22）低 48.7%，級距 −25%～−50%；Comparable 單獨為 −43.4%。主因是倍數不是 EPS：本模型 12 個月後前瞻 EPS 119.49 比共識前瞻 EPS 103.48 高 15.5%，但同業（高通、博通、邁威爾、世芯、創意）前瞻 P/E 中位數 19.0 倍，現價隱含 38.8 倍、共識目標價 6,067 隱含 58.6 倍（比同業高 208.9%）；市場給聯發科的 AI ASIC 成長溢價，Comparable 未計入。', '情境分歧大：FY2028 EPS A／B／C 為 149.3／119.5／85.3，DCF 價格帶 1,012～2,466；C 情境（顆數 ×0.6、手機 −8%）FY2028 EPS 低於 FY2027 的 100.1。即使 A 情境 DCF 2,466 也比現價低 38.5%。', 'TPU 份額是最敏感變數：聯發科占 TPU −5pt 時 FY2027 EPS −13.2%、取上限 20% 時 +13.2%；營業費用彈性 0.6 時 −12.7%（Sens 頁）。下一個檢驗點是 3Q26 財報與法說（約 2026-10 底）的 ASIC 營收與客戶進度。']   # 重點觀察（與 README.md 相同文字；數字取自重算後的 Val／PL／Sens 頁）
_rr = RD.max_row + 1
RD.cell(_rr, 1, 'v0.1.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁末新增「上檔空間級距」「投資建議」兩列（Andy 2026-10-25 14:00、14:17 固定規則），README 頁末新增投資建議與重點觀察三點；既有儲存格位置與數值不變。')
RD.cell(_rr + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"').font = Font(bold=True)
for _j, _t in enumerate(OBS, 1): RD.cell(_rr + 1 + _j, 1, f"重點觀察 {_j}：" + _t)
if os.environ.get("RATING_ROWMAP"):   # 除錯用：輸出新增列號（供撰寫 html.json）；CI 不設此變數
    print("ROWMAP", GRADE_ROW, REC_ROW, _rr, _rr + 1, _rr + 2)


order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("MTK_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE, "UN_NOTE": UN_NOTE, "T": T, "T10": T10,
               "TPS0": TPS0, "TPS_LAST": TPS_LAST, "PEER_LAST": PEER_LAST, "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0},
              open(os.environ["MTK_ROWMAP"], "w"), ensure_ascii=False, indent=1)
