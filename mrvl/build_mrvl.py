# Marvell（MRVL）收支與估值模型 v0.1（AI 半導體 Project：運算層，客製 ASIC＋光互連 DSP＋交換器；樣板 avgo/）
# 用法：python3 mrvl/build_mrvl.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度 FY(t)＝t−1 年 2 月～t 年 1 月底（FY2027：2026-02-01～2027-01-30），11/12 落在日曆年 t−1，
# 所以 FY(t) 直接取 CY(t−1)（與 NVIDIA 相同，平移一年；不跨年加權）。金額：十億美元（US$B）；EPS：美元；口徑：公司 non-GAAP（排除股票薪酬、攤銷、重組、併購費用）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
XP = ("trn", "oth", "tpu")                            # Marvell 的 XPU 來自這三個市場層平台
PNAME = {"trn": "AWS Trainium", "oth": "其他 ASIC（MTIA、Maia、Ascend 等）", "tpu": "Google TPU"}

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); NET_F = only("markets/network/dist/*.xlsx"); AVGO_F = only("avgo/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)

ws_ = openpyxl.load_workbook(TSMC_F, data_only=True)["In"]
TS = {}
for r in ws_.iter_rows(min_row=4):
    k = r[0].value
    if k and (k.startswith("um2") or k.startswith("g29_") or k.startswith("u_") or k.startswith("g_")):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
for k in XP:
    for y in (26, 27, 28): assert f"u_{k}_{y}" in TS, k
    assert f"g29_{k}" in TS

wav = openpyxl.load_workbook(AVGO_F, data_only=True)["In"]
AV = {}
for r in wav.iter_rows(min_row=4):
    k = r[0].value
    if k and (k.startswith(("core_", "sh_", "gb_", "hbm_")) or k in ("f_hbm", "m_hbm", "pa_trn")):
        AV[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(AVGO_F)}｜In!E{r[0].row}")
for k in ("core_tpu", "core_oth", "f_hbm", "m_hbm", "pa_trn", "gb_tpu_26", "gb_oth_26", "sh_tpu_28", "sh_oth_28", "hbm_A_26", "hbm_C_29"):
    assert k in AV, k

wsn = openpyxl.load_workbook(NET_F, data_only=True)
NET = {}
for r in wsn["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in ("  其中可插拔", "  其中 CPO", "scale-out 光模組"):
        NET[lab] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
for r in wsn["In"].iter_rows(min_row=4):
    if r[0].value == "rep_opt26":
        NET["rep_opt26"] = (r[4].value, f"{rel(NET_F)}｜In!E{r[0].row}")
assert "  其中可插拔" in NET and "rep_opt26" in NET
GBT = {"trn": (110, 144, 200)}   # markets/memory UNITS（與台積電模型同一份；Trainium 每顆 HBM）
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
import market as MEM
for k, nm, u_, gb, tag, src in MEM.UNITS:
    if k in XP:
        GBT[k] = gb
        for i, y in enumerate((26, 27, 28)):
            assert abs(TS[f"u_{k}_{y}"][0] - u_[i]) < 1e-9, (k, y)

PEAK, PEAK_SRC = CP.peak_years(ROOT)

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；每顆價格：千美元（US$K）；每模組金額：美元；顆數與模組數：百萬。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝Marvell 8-K 新聞稿（SEC EDGAR、investor.marvell.com）；公司對 FY2027、FY2028 的展望與投資人日目標為 Interested-party，FY2027 全年展望只用來推第四季（錨點），FY2028 起只放 Recon 對帳，不作校準"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
EDG = "https://www.sec.gov/Archives/edgar/data/1835632/"
Q4PR = EDG + "000183563226000006/q426_8kx1312026ex-991.htm"
Q1PR = "https://investor.marvell.com/sec-filings/all-sec-filings/content/0001835632-26-000014/q127_8kx522026ex-991.htm"
Q2PR = "https://investor.marvell.com/news-events/press-releases/detail/1031/marvell-technology-inc-reports-second-quarter-of-fiscal-year-2027-financial-results"
Q3PR26 = "https://investor.marvell.com/news-events/press-releases/detail/999/marvell-technology-inc-reports-third-quarter-of-fiscal-year-2026-financial-results"
CALL2 = "Q2 FY27 法說（2026-08-27）轉述：singjupost 逐字稿、Zacks 摘要（二手）"
CALL4 = "Q4 FY26 法說（2026-03-05）轉述：roic.ai 逐字稿（二手）"
IDAY = "投資人日（2026-10-06）轉述：Futurum、Benzinga、Stocktwits（二手）"

sec("A. 加速器顆數（markets/memory；與台積電、HBM、網通、NVIDIA、Broadcom 模型同一份；情境係數讀台積電模型成品，不另訂）")
for k in XP:
    for i, y in enumerate((26, 27, 28)):
        inp(f"u_{k}_{y}", f"{PNAME[k]} 顆數 CY20{y}", "百萬顆", None, TS[f"u_{k}_{y}"][0], None, "Interested-party" if k == "tpu" else ("Assumed" if k == "trn" else "Derived／Assumed"),
            TS[f"u_{k}_{y}"][1], "markets/memory/market.py UNITS（值相同，建置時檢查）" if i == 0 else "同上")
    inp(f"g29_{k}", f"{PNAME[k]} CY2029 顆數成長", "%", None, TS[f"g29_{k}"][0], None, "Assumed", TS[f"g29_{k}"][1], "markets/memory 無 2029；沿用台積電模型")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：CY2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、Broadcom 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：CY2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
for k in XP:
    for i, y in enumerate((26, 27, 28)):
        inp(f"gb_{k}_{y}", f"{PNAME[k]} 每顆 HBM CY20{y}", "GB", None, GBT[k][i], None, "Derived／Assumed", "markets/memory/market.py UNITS", "CY2029 沿用 CY2028" if i == 0 else "同上")

sec("B. ASIC 每顆規格（權威在 Broadcom 模型 Units／In 頁，唯讀；本模型只加 Marvell 份額）與 Marvell 份額")
for k, lab in (("core_tpu", "TPU 每顆核心售價（不含 HBM）"), ("core_oth", "其他 ASIC 每顆核心售價（不含 HBM）")):
    lo, b_, hi = AV[k][0]
    inp(k, lab, "US$K", lo, b_, hi, "Assumed", AV[k][1], "Broadcom 模型的 ASIC 每顆規格（本專案權威，規則第 13 條）；本模型不另訂")
lo, b_, hi = AV["f_hbm"][0]
inp("f_hbm", "HBM 由 ASIC 設計夥伴代購並計入營收的比例", "x", lo, b_, hi, "Assumed", AV["f_hbm"][1], "沿用 Broadcom 模型；Marvell 客戶是否自購 HBM：未找到")
lo, b_, hi = AV["m_hbm"][0]
inp("m_hbm", "代購 HBM 的毛利率", "x", lo, b_, hi, "Assumed", AV["m_hbm"][1], "沿用 Broadcom 模型")
inp("pa_trn", "AWS Trainium 每顆（含 HBM，模組層級）", "US$K", None, AV["pa_trn"][0][1], None, "Derived", AV["pa_trn"][1], "Broadcom 模型讀自 NVIDIA 模型 In 頁")
inp("core_trn", "Trainium 每顆核心售價（不含 HBM）＝ 每顆 − CY2026 HBM 金額", "US$K", 4.5, "=pa_trn-gb_trn_26*hbm_A_26/1000", 8.0, "Derived", "pa_trn − 110 GB × CY2026 HBM 單價", "低／高 Assumed；Broadcom 模型無 Trainium 規格")
for k in ("tpu", "oth"):
    for y in (26, 27, 28):
        inp(f"av_{k}_{y}", f"Broadcom 占 {PNAME[k]} 顆數 CY20{y}（對照；份額加總檢查用）", "%", None, AV[f"sh_{k}_{y}"][0][1], None, "Assumed", AV[f"sh_{k}_{y}"][1], "Broadcom 模型；Marvell＋Broadcom 不得超過 100%（建置時檢查）")
MSH = {("trn", 26): (0.05, 0.10, 0.20), ("trn", 27): (0.05, 0.10, 0.20), ("trn", 28): (0.0, 0.08, 0.20),
       ("oth", 26): (0.02, 0.05, 0.10), ("oth", 27): (0.03, 0.06, 0.12), ("oth", 28): (0.08, 0.15, 0.25),
       ("tpu", 26): (0.0, 0.0, 0.02), ("tpu", 27): (0.0, 0.0, 0.05), ("tpu", 28): (0.0, 0.05, 0.15)}
MSRC = {"trn": ("Marvell 供應 Trainium 2、Inferentia；Trainium 3 主供應商為世芯（Alchip），Marvell 做 Trainium 2.5（Dolphin Research 2026-03-06，二手）", "各世代份額未揭露：未找到"),
        "oth": ("Microsoft：Maia 由 Marvell 設計，Maia 300 自 CY2027 起多 GW 部署（Motley Fool／Webull 轉述 Nadella，二手）；新 Tier 1 XPU 於 FY2028 量產（公司，Q4 FY26 法說）", "其他 ASIC 含 MTIA、昇騰等非 Marvell 晶片；Broadcom 占 35%／55%／55%"),
        "tpu": ("Google：擴大合作（2026-07，客製處理器與相關晶片）；認股權證協議（歸屬上限 $120B），主要營收影響在 FY2029 起（投資人日，二手）", "Google 案是否為 TPU 平台：未找到（可能為 Axion CPU 或其他）；Broadcom 占 TPU 75%（CY2028）")}
for k in XP:
    for y in (26, 27, 28):
        lo, b_, hi = MSH[(k, y)]
        inp(f"msh_{k}_{y}", f"Marvell 占 {PNAME[k]} 顆數 CY20{y}" + ("（CY2029 沿用）" if y == 28 else ""), "%", lo, b_, hi, "Assumed", MSRC[k][0] if y == 26 else "同上", MSRC[k][1] if y == 26 else "—")
        if k != "trn":
            assert AV[f"sh_{k}_{y}"][0][1] + hi <= 1.0 + 1e-9, (k, y)    # Broadcom＋Marvell（高）≤ 100%
inp("pfc", "CY2027 起 XPU 核心售價係數（Sens 用；基準 1）", "x", 0.8, 1.0, 1.2, "Assumed", "無：新世代核心單價變化未揭露", "Sens 用 ±20%")
inp("uk", "Marvell XPU 顆數係數（Sens 用；CY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "無", "±20%")

sec("C. 記憶體價格路徑（讀 Broadcom 模型成品，其來源為 skhynix 成品；唯讀；情境別）")
for s in SC:
    for y in (26, 27, 28, 29):
        inp(f"hbm_{s}_{y}", f"HBM 混合單價 CY20{y}（{s}）", "US$/GB", None, AV[f"hbm_{s}_{y}"][0][1], None, "Derived", AV[f"hbm_{s}_{y}"][1], "markets/memory：2026 BofA $14.6/GB；2027 +80%" if (s == "A" and y == 26) else "同上")

sec("D. 光互連：AI 後端可插拔光模組（讀 markets/network 成品；情境 A 基準，B／C 依顆數係數縮放）× 每模組晶片組金額 × Marvell 份額")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"plug_{y}", f"AI 後端可插拔光模組市場 CY20{y}", "US$B", None, NET["  其中可插拔"][0][j], None, "Derived", NET["  其中可插拔"][1], "markets/network：scale-out 光模組扣 CPO；2026 較 LightCounting 26（含 scale-up）高約 16%" if y == 26 else "同上")
inp("rep_opt26", "AI 叢集光學市場 2026（LightCounting；對照）", "US$B", None, NET["rep_opt26"][0], None, "Interested-party", NET["rep_opt26"][1], "—")
for y, (lo, b_, hi) in zip((26, 27, 28, 29), ((550, 650, 750), (650, 800, 950), (700, 850, 1000), (700, 850, 1000))):
    inp(f"asp_{y}", f"光模組平均單價 CY20{y}", "US$", lo, b_, hi, "Assumed", "無：800G 約 $500–600、1.6T 約 $1,000 的組合（市場轉述，二手）；1.6T 占比上升" if y == 26 else "同上", "市場別單價：未找到")
for y, (lo, b_, hi) in zip((26, 27, 28, 29), ((0.0, 0.03, 0.08), (0.02, 0.06, 0.15), (0.04, 0.10, 0.25), (0.06, 0.15, 0.30))):
    inp(f"lpo_{y}", f"不用 DSP 的線性直驅（LPO／LRO）占模組數 CY20{y}", "%", lo, b_, hi, "Assumed", "LightCounting 2026-04：線性直驅 2027–2031 年壓抑 DSP 銷售（未量化，Interested-party）" if y == 26 else "同上", "第二來源：未找到")
for y, (lo, b_, hi) in zip((26, 27, 28, 29), ((100, 125, 160), (110, 135, 170), (110, 135, 170), (105, 130, 165))):
    inp(f"dsp_{y}", f"每模組光晶片組金額（PAM4 DSP＋TIA／驅動器）CY20{y}", "US$", lo, b_, hi, "Assumed", "LightCounting：2025 年 800G PAM4 晶片組出貨近三倍、2026 年 1.6T 晶片組 >$2B（Interested-party）" if y == 26 else "同上", "約為模組單價 15–20%（Assumed）；第二來源：未找到")
for y, (lo, b_, hi) in zip((26, 27, 28, 29), ((0.50, 0.65, 0.75), (0.45, 0.63, 0.72), (0.40, 0.60, 0.70), (0.35, 0.58, 0.70))):
    inp(f"msh_dsp_{y}", f"Marvell 占光晶片組金額 CY20{y}", "%", lo, b_, hi, "Assumed", "PAM4 DSP 主要供應商：Marvell、Broadcom、Credo、MaxLinear；Marvell 份額未揭露" if y == 26 else "同上", "第二來源：未找到（LightCounting 份額為付費資料）")
inp("coh27", "相干 DSP／DCI 與前端光學（非 AI 後端）FY2027", "US$B", 0.6, 0.9, 1.2, "Assumed", "公司：1.6T ZR／ZR+ 相干 DCI 模組（Q2 FY27 法說，二手）；LightCounting：2025 年相干 DSP +16%", "金額未揭露：未找到")
inp("g_coh", "相干 DSP／DCI 年增（FY2028 起）", "%", 0.05, 0.15, 0.25, "Assumed", "無", "—")
inp("of", "Marvell 光互連係數（Sens 用；FY2028 起）", "x", 0.8, 1.0, 1.2, "Assumed", "無", "±20%")

sec("E. 資料中心其他產品線（公司說法；Interested-party）")
inp("cus26", "客製（XPU＋XPU attach）營收 FY2026", "US$B", None, 1.5, None, "Interested-party", CALL4 + "：「從零成長到 FY2026 的 $1.5B」；年增一倍", "第二來源：未找到")
inp("g_cus27", "客製營收年增 FY2027", "%", 0.20, 0.25, 0.35, "Interested-party", "Q4 FY26、Q1 FY27 法說：「超過 20%，偏高」；Q2 FY27：下半年顯著加速（二手）", "共識無拆分：未找到")
inp("att26", "XPU attach（客製 NIC、CXL 等）FY2026", "US$B", 0.15, 0.20, 0.30, "Interested-party", CALL4 + "：「去年約數億美元，今年倍增」", "—")
inp("att27", "XPU attach FY2027（倍增）", "US$B", 0.3, 0.4, 0.5, "Interested-party", "同上", "—")
inp("att28", "XPU attach FY2028", "US$B", 0.6, 1.0, 1.4, "Interested-party", CALL4 + "：「明年大概是十億美元級的業務」；FY2028 超過 10 個 attach 案放量", "—")
inp("g_att", "XPU attach 年增（FY2029 起）", "%", 0.3, 0.5, 0.8, "Assumed", "投資人日：FY2031 客製約 $30B，XPU 與 attach 大致各半（二手）", "—")
inp("sw26", "資料中心交換器（Teralynx）FY2026", "US$B", None, 0.30, None, "Interested-party", CALL4 + "：超過 $300M", "Futurum（Q1 FY27）：FY2026 約 $300M 隱含")
inp("sw27", "資料中心交換器 FY2027", "US$B", 0.6, 0.65, 0.75, "Interested-party", CALL4 + "：超過 $600M；Q2 FY27：今年倍增以上", "Zacks（二手）同")
inp("g_sw28", "交換器年增 FY2028", "%", 0.4, 0.8, 1.2, "Interested-party", "Q2 FY27 法說：朝 $1B 年化營收（二手）", "—")
inp("g_sw", "交換器年增 FY2029 起", "%", 0.2, 0.4, 0.6, "Assumed", "UALink、scale-up 交換器開發中", "—")
inp("aec26", "AEC 與 retimer FY2026", "US$B", 0.15, 0.20, 0.25, "Interested-party", CALL4 + "：「大概 $200M 範圍」", "—")
inp("aec27", "AEC 與 retimer FY2027（倍增以上）", "US$B", 0.4, 0.45, 0.55, "Interested-party", "同上：「FY2027 倍增以上」", "—")
inp("g_aec", "AEC 與 retimer 年增 FY2028 起", "%", 0.2, 0.5, 0.8, "Assumed", "無", "—")
inp("sto26", "儲存控制器與其他資料中心 FY2026", "US$B", 0.5, 0.65, 0.8, "Assumed", "無：公司不揭露；Dolphin Research：資料中心含客製、光互連、SSD 控制器（二手）", "第二來源：未找到")
inp("g_sto", "儲存與其他年增", "%", -0.05, 0.05, 0.15, "Assumed", "無", "—")

sec("F. FY2027 錨點（Q1、Q2 實績＋Q3 指引中點＋Q4 由公司全年展望推得）")
inp("q3_rev", "Q3 FY27 營收指引", "US$B", None, 3.15, None, "Interested-party", Q2PR + "：$3,150M ± 5%", "Zacks、Futurum（二手）同值")
inp("q3_cog", "Q3 FY27 通訊與其他年增（指引說法）", "%", -0.15, -0.135, -0.11, "Interested-party", CALL2 + "：季減與年減皆為「十幾個百分點低至中段」", "—")
inp("q3_gm", "Q3 FY27 non-GAAP 毛利率指引中點", "%", 0.575, 0.58, 0.585, "Interested-party", Q2PR + "：57.5%–58.5%", "—")
inp("q3_opx", "Q3 FY27 non-GAAP 營業費用指引", "US$B", None, 0.655, None, "Interested-party", Q2PR + "：約 $655M", "—")
inp("q3_oth", "Q3 FY27 non-GAAP 利息與其他淨損", "US$B", 0.030, 0.036, 0.045, "Interested-party", CALL2 + "：約 $36M", "Q1 0.040、Q2 0.030（公司）")
inp("q3_tx", "Q3 FY27 non-GAAP 稅率", "%", None, 0.11, None, "Interested-party", CALL2 + "：11%", "Q1、Q2 11.0%（公司）")
inp("q3_sh", "Q3 FY27 稀釋股數指引", "十億股", None, 0.921, None, "Interested-party", Q2PR + "：921M", "—")
inp("co_rev27", "公司 FY2027 全年營收展望", "US$B", 11.8, 12.0, 12.2, "Interested-party", CALL2 + "：約 $12B（年增約 45%）", "投資人日（Benzinga，二手）：仍約 $12B；共識 12.07 只放 Recon")
inp("co_dcg27", "公司 FY2027 資料中心營收年增展望", "%", 0.55, 0.60, 0.65, "Interested-party", CALL2 + "：約 60%（原約 50%）", "Q3 資料中心季增 >20%、年增約 75%（同上）")
inp("q4_gm", "Q4 FY27 non-GAAP 毛利率", "%", 0.57, 0.58, 0.59, "Interested-party", CALL2 + "：Q4 與 Q3 同區間", "—")
inp("co_opx27", "公司 FY2027 non-GAAP 營業費用", "US$B", 2.50, 2.55, 2.60, "Interested-party", CALL2 + "：約 $2.55B（原 $2.45B）", "—")
inp("q4_oth", "Q4 FY27 non-GAAP 利息與其他淨損", "US$B", 0.030, 0.036, 0.045, "Assumed", "沿用 Q3 指引", "—")
inp("q4_tx", "Q4 FY27 non-GAAP 稅率", "%", 0.10, 0.11, 0.12, "Assumed", "沿用 Q3", "—")
inp("q4_sh", "Q4 FY27 稀釋股數", "十億股", 0.915, 0.925, 0.935, "Assumed", "Q3 指引 0.921 ＋ 股票薪酬稀釋", "—")

sec("G. 營業費用、利息、稅、股數（FY2028 起）、通訊與其他")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.3, 0.5, 0.7, "Interested-party", CALL2 + "：FY2028 營業費用成長約為營收成長的一半", "FY2027 錨點約 0.27（Q 頁）")
inp("g_opx_min", "營業費用年增下限", "%", 0.0, 0.05, 0.10, "Assumed", "無", "—")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 ±2 個百分點")
inp("int_f", "non-GAAP 利息與其他淨損（每年）", "US$B", 0.10, 0.15, 0.30, "Derived", "FY2026 0.151、FY2027 錨點約 0.14（Q 頁）；總債務 $4.96B（Q2 FY27 法說，二手）", "高端：產能預付約 $1B 以舉債支應（Assumed）")
inp("tx_f", "non-GAAP 稅率（FY2028 起）", "%", 0.11, 0.13, 0.15, "Interested-party", CALL2 + "：FY2028 約 13%", "FY2026 10%、FY2027 11%（公司）")
inp("sh_g", "稀釋股數年變動", "%", -0.01, 0.005, 0.015, "Derived", "FY2026 0.870 → Q3 FY27 0.921（含 Celestial AI 併購發股，公司）", "低／高 Assumed；買回 Q2 $200M")
inp("g_co28", "通訊與其他年增 FY2028", "%", -0.05, 0.03, 0.08, "Interested-party", "Dolphin Research 轉述公司：FY2028 低個位數（二手）", "第二來源：未找到")
inp("g_co", "通訊與其他年增 FY2029 起", "%", -0.05, 0.03, 0.08, "Assumed", "企業網通、電信、消費、車用／工業；車用乙太網路已於 2025-08 售予 Infineon", "—")
inp("gm_opt", "光互連毛利率", "x", 0.58, 0.63, 0.68, "Assumed", "無：公司不揭露產品別毛利率", "第二來源：未找到")
inp("gm_oth", "交換器、AEC、儲存與其他資料中心毛利率", "x", 0.55, 0.60, 0.65, "Assumed", "無", "—")
inp("gm_att", "XPU attach 毛利率", "x", 0.45, 0.55, 0.65, "Assumed", "無", "—")
inp("gm_co", "通訊與其他毛利率", "x", 0.53, 0.58, 0.63, "Assumed", "無：FY2026 全公司 59.5%（公司）", "—")

sec("H. 公司展望、投資人日與共識（只作 Recon 對帳，不作校準）")
inp("co_rev28", "公司 FY2028 營收展望（投資人日）", "US$B", None, 20.0, None, "Interested-party", IDAY + "：約 $20B（8 月為 $18B、3 月為 $15B）", "Stocktwits：資料中心約 $18B")
inp("co_dc28", "公司 FY2028 資料中心營收（投資人日）", "US$B", None, 18.0, None, "Interested-party", "Stocktwits 轉述投資人日（二手，單一來源）", "—")
inp("co_cus29", "公司 FY2029 客製營收目標（下限）", "US$B", None, 12.0, None, "Interested-party", IDAY + "：超過 $12B（原超過 $10B）；為 FY2028 的三倍以上", "—")
inp("co_icg27", "公司 FY2027 光互連（interconnect）年增（下限）", "%", None, 0.70, None, "Interested-party", "Q1 FY27 法說：超過 70%（Futurum，二手）", "Q4 FY26：超過 50%")
inp("co_eps28", "公司 FY2028 non-GAAP EPS（3 月說法：下限）", "US$", None, 5.0, None, "Interested-party", CALL4 + "：「遠高於 $5」（當時營收展望 $15B）", "投資人日 EPS：未找到")
inp("co_gm31", "公司 FY2031 non-GAAP 毛利率目標（中點）", "%", 0.56, 0.575, 0.59, "Interested-party", IDAY + "：56%–59%", "—")
inp("cons_rev27", "共識 FY2027 營收", "US$B", None, 12.07, None, "Interested-party", "StockAnalysis（S&P Global，41 位分析師；2026-10-07 更新）", "—")
inp("cons_eps27", "共識 FY2027 non-GAAP EPS", "US$", None, 4.22, None, "Interested-party", "同上（區間 3.92–4.61）", "—")
inp("cons_rev28", "共識 FY2028 營收", "US$B", None, 19.71, None, "Interested-party", "同上（「明年」欄）", "Benzinga：投資人日前共識 $18.2B")
inp("cons_eps28", "共識 FY2028 non-GAAP EPS", "US$", None, 7.33, None, "Interested-party", "同上", "—")

sec("I. 價格、估值與目標價權重")
inp("px", "股價（2026-10-23 收盤）", "US$", None, 228.51, None, "Interested-party", "StockAnalysis（2026-10-23 收盤）", "Benzinga 報價 227.70（盤後）；10-06 投資人日後 $282，10-09～10-16 跌 17%（Trefis）")
inp("pe_ttm_gaap", "近 12 月 GAAP P/E（2026-10-16）", "x", None, 77.4, None, "Interested-party", "Trefis 2026-10-19", "第二來源：未找到")
inp("eps_q3_26", "Q3 FY26 non-GAAP EPS（近 12 月計算用）", "US$", None, 0.76, None, "Verified", Q3PR26, "—")
inp("eps_q4_26", "Q4 FY26 non-GAAP EPS（近 12 月計算用）", "US$", None, 0.80, None, "Verified", Q4PR, "—")
inp("tp_wA", "目標價情境權重 A", "x", None, 0.3, None, "Assumed", "Andy 可改", "權重合計應為 1（Val 頁檢查）")
inp("tp_wB", "目標價情境權重 B", "x", None, 0.5, None, "Assumed", "Andy 可改", "—")
inp("tp_wC", "目標價情境權重 C", "x", None, 0.2, None, "Assumed", "Andy 可改", "—")
for s in SC:
    inp(f"pk_{s}", f"情境 {s}：設備市場見頂年度（日曆年；循環位置規則用）", "年", None, PEAK[s], None, "Derived", PEAK_SRC[s], "共用規則（tools/cycle_position.py）：Marvell 的 FY(t) 對照日曆年 t−1，用 classify_lag（與 NVIDIA 相同）" if s == "A" else "同上")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2027 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（十億美元；EPS 為美元；股數為十億股）。口徑：公司 non-GAAP（排除股票薪酬、攤銷、重組、併購費用）。藍字＝輸入；黑字＝公式。實績取自 Marvell 8-K 新聞稿（SEC EDGAR、investor.marvell.com）"
Q["A2"] = "FY2027 錨點＝Q1、Q2 實績＋Q3 指引中點＋Q4（＝公司全年營收展望約 $12B − 前三季；資料中心依全年年增約 60% 推得）；假設見 In 頁 F 節。資料中心產品線拆分（客製、交換器、AEC、儲存）取公司法說說法，光互連為差額（Derived）"
hdr(Q, 3, ["項目", "FY2026A（～2026-01-31）", "FY27 Q1（～2026-05-02）", "FY27 Q2（～08-01）", "FY27 Q3（指引）", "FY27 Q4（推得）", "FY2027 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = "FY2026：" + Q4PR + "；Q1 FY27：" + Q1PR + "；Q2 FY27：" + Q2PR
S2 = "交叉：Q2 FY27 稿的六個月合計（營收 5.1571、non-GAAP 營業利益 1.8501、non-GAAP 淨利 1.5839）＝Q1＋Q2（下方檢查列）；FY2026 資料中心＝四季加總（Q1 1.4406、Q2 1.4905、Q3 1.5179、Q4 1.6513）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
qrow(4, "營收", {"B": 8.1946, "C": 2.4178, "D": 2.7393, "E": "=q3_rev", "F": "=co_rev27-C4-D4-E4"}, tag="Verified／Interested-party（Q3 指引、全年展望）")
qrow(5, "  資料中心", {"B": "=1.4406+1.4905+1.5179+1.6513", "C": 1.8327, "D": 2.1715, "E": "=E4-E6", "F": "=B5*(1+co_dcg27)-C5-D5-E5"}, tag="Verified／Derived（Q3、Q4）", s2="FY2026 四季：Q1 FY27 稿（Q1 FY26 比較欄）、Q3 FY26 稿（Q2、Q3）、Q4 FY26 稿（Q4）")
qrow(6, "  通訊與其他（企業網通、電信、消費、車用／工業）", {"B": "=B4-B5", "C": 0.5851, "D": 0.5678, "E": "=(2.0745-1.5179)*(1+q3_cog)", "F": "=F4-F5"}, tag="Verified／Derived", s2="Q3 FY26 通訊與其他＝營收 2.0745 − 資料中心 1.5179；Q4 FY26 起公司只揭露兩類")
qrow(7, "  資料中心：客製（XPU＋XPU attach）", {"B": "=cus26"}, ann="=cus26*(1+g_cus27)", tag="Interested-party", s1=CALL4, s2="FY2027＝FY2026 × (1＋g_cus27)")
qrow(8, "    其中 XPU attach", {"B": "=att26"}, ann="=att27", tag="Interested-party", s1=CALL4, s2="—")
qrow(9, "    其中 XPU（＝客製 − attach）", {"B": "=B7-B8"}, ann="=G7-G8", tag="Derived", s1="—", s2="—")
qrow(10, "  資料中心：交換器", {"B": "=sw26"}, ann="=sw27", tag="Interested-party", s1=CALL4, s2="—")
qrow(11, "  資料中心：AEC 與 retimer", {"B": "=aec26"}, ann="=aec27", tag="Interested-party", s1=CALL4, s2="—")
qrow(12, "  資料中心：儲存與其他", {"B": "=sto26"}, ann="=sto26*(1+g_sto)", tag="Assumed", s1="—", s2="—")
qrow(13, "  資料中心：光互連（＝資料中心 − 上列；Derived）", {"B": "=B5-B7-B10-B11-B12"}, ann="=G5-G7-G10-G11-G12", tag="Derived", s1="—", s2="含 PAM4 DSP、TIA／驅動器、相干 DSP／DCI、scale-up 光學")
qrow(14, "    光互連年增（對照公司「超過 70%」）", {}, fmt="0.0%", ann="=G13/B13-1", tag="Derived", s1="—", s2="公司 FY2027 說法：超過 70%（Q1 FY27 法說）")
qrow(15, "non-GAAP 毛利", {"B": 4.8718, "C": 1.4238, "D": 1.6140, "E": "=E4*q3_gm", "F": "=F4*q4_gm"}, tag="Verified／Interested-party（Q3、Q4）")
qrow(16, "毛利率", {c: f"={c}15/{c}4" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2026 59.5%；Q3 指引 57.5–58.5%")
qrow(17, "non-GAAP 營業費用", {"B": 1.9808, "C": 0.5769, "D": 0.6108, "E": "=q3_opx", "F": "=co_opx27-C17-D17-E17"}, tag="Verified／Interested-party")
qrow(18, "non-GAAP 營業利益", {c: f"={c}15-{c}17" for c in "BCDEF"}, tag="Derived", s2="公司揭露：FY2026 2.8910、Q1 0.8469、Q2 1.0032")
qrow(19, "營業利益率", {c: f"={c}18/{c}4" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="公司：Q4 FY27 進入 38–40% 目標區間")
qrow(20, "non-GAAP 利息與其他淨損", {"B": 0.1514, "C": 0.0401, "D": 0.0303, "E": "=q3_oth", "F": "=q4_oth"}, tag="Verified／Interested-party")
qrow(21, "稅前利益", {c: f"={c}18-{c}20" for c in "BCDEF"}, tag="Derived")
qrow(22, "non-GAAP 稅率", {"B": 0.10, "C": 0.11, "D": 0.11, "E": "=q3_tx", "F": "=q4_tx"}, fmt="0.0%", ann="=G23/G21", tag="Verified／Interested-party")
qrow(23, "non-GAAP 所得稅", {c: f"={c}21*{c}22" for c in "BCDEF"}, tag="Derived")
qrow(24, "non-GAAP 淨利", {c: f"={c}21-{c}23" for c in "BCDEF"}, tag="Derived")
qrow(25, "  公司揭露 non-GAAP 淨利", {"B": 2.4656, "C": 0.7180, "D": 0.8659}, ann="", tag="Verified")
qrow(26, "non-GAAP 稀釋股數（十億股）", {"B": 0.8697, "C": 0.8933, "D": 0.9212, "E": "=q3_sh", "F": "=q4_sh"}, ann="=AVERAGE(C26:F26)", tag="Verified／Interested-party（Q3）／Assumed（Q4）")
qrow(27, "non-GAAP EPS", {c: f"={c}24/{c}26" for c in "BCDEF"}, fmt="0.000", ann="=G24/G26", tag="Derived", s1="淨利 ÷ 稀釋股數", s2="—")
qrow(28, "  公司揭露 non-GAAP EPS", {"B": 2.84, "C": 0.80, "D": 0.94}, fmt="0.00", ann="", tag="Verified")
qrow(29, "檢查：淨利 − 公司揭露（應約為 0；四捨五入）", {c: f"={c}24-{c}25" for c in "BCD"}, fmt="0.000", ann="", tag="—", s1="—", s2="—")
qrow(30, "檢查：Q1＋Q2 − 六個月合計（營收、營業利益、淨利；應約為 0）", {"G": "=(C4+D4-5.1571)+(C18+D18-1.8501)+(C24+D24-1.5839)"}, fmt="0.000", ann="", tag="—", s1=Q2PR + "（六個月欄）", s2="—")
Q.cell(32, 1, "FY2027 錨點與共識、公司展望的距離（參考；校準未使用共識）").font = BOLD
Q.cell(33, 1, "FY2027 營收 ÷ 共識 − 1"); Q["B33"] = "=G4/cons_rev27-1"
Q.cell(34, 1, "FY2027 EPS ÷ 共識 − 1"); Q["B34"] = "=G27/cons_eps27-1"
Q.cell(35, 1, "FY2027 資料中心年增"); Q["B35"] = "=G5/B5-1"
Q.cell(36, 1, "FY2027 營收年增"); Q["B36"] = "=G4/B4-1"
Q.cell(37, 1, "FY2027 營業費用年增"); Q["B37"] = "=G17/B17-1"
Q.cell(38, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；FY2027 錨點）"); Q["B38"] = "=B37/B36"
for rr in range(33, 39): Q[f"B{rr}"].number_format = "0.0%" if rr < 38 else "0.00"
for nm, rr in (("rev", 4), ("dc", 5), ("co", 6), ("cus", 7), ("att", 8), ("xpu", 9), ("sw", 10), ("aec", 11), ("sto", 12), ("ic", 13),
               ("gp", 15), ("opx", 17), ("oi", 18), ("oth", 20), ("tax", 23), ("ni", 24), ("sh", 26), ("eps", 27)):
    defname(f"a26_{nm}", f"Q!$B${rr}"); defname(f"a27_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEFG": Q.column_dimensions[c].width = 15
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 90; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
FYS = (2026, 2027, 2028, 2029, 2030); FC = {2026: "C", 2027: "D", 2028: "E", 2029: "F", 2030: "G"}
BU = (2027, 2028, 2029, 2030)          # 有由下而上推導的會計年度；FY(t)≈CY(t−1)；FY2027 用來校準
cy = lambda fy: fy - 1
def uexp(k, c, s):
    if c <= 2027: return f"u_{k}_{c % 100}"
    if c == 2028: return f"u_{k}_28*um28_{s}"
    return f"u_{k}_28*(1+g29_{k})*um29_{s}"
umx = lambda c, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[c]
DEFP = dict(scn="A", uf="1", sht="0", sho="0", shp="0", pfc="1", fh="f_hbm", mh="m_hbm", of="1", dspf="1", lpoa="0", shd="0",
            a28="att28", gsw28="g_sw28", gco28="g_co28", opxa="0", flat=False)
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "FY2030", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    yc = lambda y: min(cy(y), 2028) if (P["flat"] and cy(y) == 2029) else cy(y)
    sec_("A. XPU（由下而上；D 欄＝FY2027≈CY2026 校準年，E–G＝FY2028–FY2030≈CY2027–CY2029）")
    for k in XP:
        row("u_" + k, f"{PNAME[k]} 市場顆數（CY）", "百萬顆", {c_(y): "=" + uexp(k, yc(y), s) for y in BU})
    shx = {"trn": P["sht"], "oth": P["sho"], "tpu": P["shp"]}
    for k in XP:
        row("s_" + k, f"Marvell 份額：{PNAME[k]}", "%", {c_(y): (f"=msh_{k}_26" if y == 2027 else f"=MIN(1,MAX(0,msh_{k}_{min(cy(y), 2028) % 100}+{shx[k]}))") for y in BU}, fmt="0.0%")
    for k in XP:
        row("b_" + k, f"Marvell 顆數：{PNAME[k]}", "百萬顆", {c_(y): f"={c_(y)}{R['u_' + k]}*{c_(y)}{R['s_' + k]}" + ("" if y == 2027 else f"*{P['uf']}") for y in BU}, fmt="0.000")
    for k in XP:
        row("g_" + k, f"每顆 HBM：{PNAME[k]}", "GB", {c_(y): f"=gb_{k}_{min(cy(y), 2028) % 100}" for y in BU}, fmt="0")
    row("hpx", "HBM 單價", "US$/GB", {c_(y): f"=hbm_{s}_{cy(y) % 100}" for y in BU})
    for k in XP:
        row("c_" + k, f"每顆核心售價：{PNAME[k]}", "US$K", {c_(y): f"=core_{k}" + ("" if y == 2027 else f"*{P['pfc']}") for y in BU})
    for k in XP:
        row("h_" + k, f"每顆 HBM 金額（全部）：{PNAME[k]}", "US$K", {c_(y): f"={c_(y)}{R['g_' + k]}*{c_(y)}{R['hpx']}/1000" for y in BU})
    for k in XP:
        row("a_" + k, f"每顆 Marvell 售價（核心＋代購 HBM）：{PNAME[k]}", "US$K", {c_(y): f"={c_(y)}{R['c_' + k]}+{P['fh']}*{c_(y)}{R['h_' + k]}" for y in BU},
            note="核心售價 ＋ 代購比例 × HBM 金額（未乘 k_x）" if k == "trn" else "")
    row("xraw", "XPU 營收（未校準）", "US$B", {c_(y): "=" + "+".join(f"{c_(y)}{R['b_' + k]}*{c_(y)}{R['a_' + k]}" for k in XP) for y in BU})
    row("kx", "XPU 校準係數 k_x ＝ FY2027 錨點 XPU ÷ 未校準", "x", {"D": f"=a27_xpu/D{r - 1}"}, fmt="0.000", note="Derived：以 FY2027 錨點（公司 FY2026 客製 $1.5B × 年增 − attach）校準；FY2028 起沿用。顆數、份額、核心單價的合併係數")
    row("xpu", "XPU 營收", "US$B", {c_(y): f"={c_(y)}{R['xraw']}*$D${R['kx']}" for y in BU})
    row("xh", "  其中代購 HBM", "US$B", {c_(y): f"=$D${R['kx']}*{P['fh']}*(" + "+".join(f"{c_(y)}{R['b_' + k]}*{c_(y)}{R['h_' + k]}" for k in XP) + ")" for y in BU})
    row("xc", "  其中核心（設計＋邏輯晶片＋封裝）", "US$B", {c_(y): f"={c_(y)}{R['xpu']}-{c_(y)}{R['xh']}" for y in BU})
    row("xu", "Marvell XPU 顆數（未乘 k_x）", "百萬顆", {c_(y): "=" + "+".join(f"{c_(y)}{R['b_' + k]}" for k in XP) for y in BU}, fmt="0.000")
    row("att", "XPU attach（客製 NIC、CXL 等）", "US$B", {"C": "=a26_att", "D": "=a27_att", "E": f"={P['a28']}", "F": f"={P['a28']}*(1+g_att)*{umx(yc(2029), s)}", "G": f"={P['a28']}*(1+g_att)^2*{umx(yc(2030), s)}"}, note="FY2029 起乘情境顆數係數")
    row("cus", "客製合計（XPU＋attach）", "US$B", {"C": "=a26_cus", **{c_(y): f"={c_(y)}{R['xpu']}+{c_(y)}{R['att']}" for y in BU}})
    sec_("B. 光互連（AI 後端可插拔光模組 × 每模組晶片組 × Marvell 份額 × k_o；另加相干 DSP／DCI）")
    row("plug", "AI 後端可插拔光模組市場（markets/network × 情境係數）", "US$B", {c_(y): f"=plug_{yc(y) % 100}*{umx(yc(y), s)}" for y in BU})
    row("mod", "光模組數量 ＝ 市場 ÷ 平均單價", "百萬個", {c_(y): f"={c_(y)}{R['plug']}*1000/asp_{yc(y) % 100}" for y in BU}, fmt="0.0")
    row("lpo", "不用 DSP 的線性直驅占比", "%", {c_(y): f"=MIN(1,lpo_{yc(y) % 100}" + ("" if y == 2027 else f"+{P['lpoa']}") + ")" for y in BU}, fmt="0.0%")
    row("dsp", "每模組光晶片組金額", "US$", {c_(y): f"=dsp_{yc(y) % 100}" + ("" if y == 2027 else f"*{P['dspf']}") for y in BU}, fmt="0")
    row("cs", "光晶片組市場 ＝ 模組數 ×（1 − 線性直驅）× 每模組金額", "US$B", {c_(y): f"={c_(y)}{R['mod']}*(1-{c_(y)}{R['lpo']})*{c_(y)}{R['dsp']}/1000" for y in BU})
    row("sd", "Marvell 份額：光晶片組", "%", {c_(y): f"=msh_dsp_{yc(y) % 100}" + ("" if y == 2027 else f"+{P['shd']}") for y in BU}, fmt="0.0%")
    row("oraw", "Marvell AI 光晶片組（未校準）", "US$B", {c_(y): f"={c_(y)}{R['cs']}*{c_(y)}{R['sd']}" + ("" if y == 2027 else f"*{P['of']}") for y in BU})
    row("coh", "相干 DSP／DCI 與前端光學", "US$B", {"D": "=coh27", **{c_(y): f"={c_(y - 1)}{r}*(1+g_coh)" for y in BU[1:]}})
    row("ko", "光互連校準係數 k_o ＝（錨點光互連 − 相干）÷ 未校準（FY2027）", "x", {"D": f"=(a27_ic-D{r - 1})/D{r - 2}"}, fmt="0.000",
        note="Derived：錨點光互連為資料中心差額；k_o 吸收 scale-up 光學、前端資料中心光學與份額、單價的合併誤差；FY2028 起沿用")
    row("ic", "光互連營收", "US$B", {"C": "=a26_ic", **{c_(y): f"={c_(y)}{R['oraw']}*$D${R['ko']}+{c_(y)}{R['coh']}" for y in BU}})
    sec_("C. 資料中心其他與通訊（FY2029 起 AI 相關產品線乘情境顆數係數）")
    gsw = P["gsw28"]
    row("sw", "交換器", "US$B", {"C": "=a26_sw", "D": "=a27_sw", "E": f"=a27_sw*(1+{gsw})", "F": f"=a27_sw*(1+{gsw})*(1+g_sw)*{umx(yc(2029), s)}", "G": f"=a27_sw*(1+{gsw})*(1+g_sw)^2*{umx(yc(2030), s)}"})
    row("aec", "AEC 與 retimer", "US$B", {"C": "=a26_aec", "D": "=a27_aec", "E": "=a27_aec*(1+g_aec)", "F": f"=a27_aec*(1+g_aec)^2*{umx(yc(2029), s)}", "G": f"=a27_aec*(1+g_aec)^3*{umx(yc(2030), s)}"})
    row("sto", "儲存與其他資料中心", "US$B", {"C": "=a26_sto", "D": "=a27_sto", **{c_(y): f"={c_(y - 1)}{r}*(1+g_sto)" for y in BU[1:]}})
    row("dc", "資料中心營收", "US$B", {"C": "=a26_dc", **{c_(y): f"={c_(y)}{R['cus']}+{c_(y)}{R['ic']}+{c_(y)}{R['sw']}+{c_(y)}{R['aec']}+{c_(y)}{R['sto']}" for y in BU}})
    row("co", "通訊與其他", "US$B", {"C": "=a26_co", "D": "=a27_co", "E": f"=D{r}*(1+{P['gco28']})", **{c_(y): f"={c_(y - 1)}{r}*(1+g_co)" for y in BU[2:]}})
    sec_("D. 報告值（FY2026A、FY2027 為實績與錨點；FY2028 起為由下而上）")
    row("rev", "營收", "US$B", {"C": "=a26_rev", "D": "=a27_rev", **{c_(y): f"={c_(y)}{R['dc']}+{c_(y)}{R['co']}" for y in BU[1:]}})
    row("bur", "  由下而上營收（FY2027 應等於錨點）", "US$B", {c_(y): f"={c_(y)}{R['dc']}+{c_(y)}{R['co']}" for y in BU})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("dcsh", "資料中心 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['dc']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("gpo", "毛利：光互連＋其他資料中心＋attach＋通訊＋代購 HBM", "US$B", {c_(y): f"=gm_opt*{c_(y)}{R['ic']}+gm_oth*({c_(y)}{R['sw']}+{c_(y)}{R['aec']}+{c_(y)}{R['sto']})+gm_att*{c_(y)}{R['att']}+gm_co*{c_(y)}{R['co']}+{P['mh']}*{c_(y)}{R['xh']}" for y in BU})
    row("mc", "XPU 核心毛利率 m_core ＝（錨點毛利 − 其他毛利）÷ 核心營收（FY2027）", "x", {"D": f"=(a27_gp-D{r - 1})/D{R['xc']}"}, fmt="0.000", note="Derived：讓 FY2027 毛利率等於錨點；FY2028 起毛利率變化來自營收組合（客製占比上升、代購 HBM）")
    row("gp", "毛利", "US$B", {"C": "=a26_gp", "D": "=a27_gp", **{c_(y): f"={c_(y)}{R['gpo']}+$D${R['mc']}*{c_(y)}{R['xc']}" for y in BU[1:]}})
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("cush", "客製 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['cus']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a26_opx", "D": "=a27_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(g_opx_min,phi*{c_(y)}{R['g']}))" for y in BU[1:]}})
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}})
    row("oi", "營業利益", "US$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("int", "利息與其他淨損", "US$B", {"C": "=a26_oth", "D": "=a27_oth", **{c_(y): "=int_f" for y in BU[1:]}})
    row("etr", "有效稅率", "%", {"C": "=a26_tax/(a26_oi-a26_oth)", "D": "=a27_tax/(a27_oi-a27_oth)", **{c_(y): "=tx_f" for y in BU[1:]}}, fmt="0.0%")
    row("ni", "淨利（non-GAAP）", "US$B", {"C": "=a26_ni", "D": "=a27_ni", **{c_(y): f"=({c_(y)}{R['oi']}-{c_(y)}{R['int']})*(1-{c_(y)}{R['etr']})" for y in BU[1:]}})
    row("sh", "稀釋股數", "十億股", {"C": "=a26_sh", "D": "=a27_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.000")
    row("eps", "EPS（non-GAAP）", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL：三情境 ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Marvell 損益（US$B；non-GAAP）。FY2026A 實績、FY2027 錨點（Q1、Q2 實績＋Q3 指引＋Q4 由全年展望推得），FY2028 起為由下而上。FY(t)≈CY(t−1)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2027–FY2028 三情境相同（顆數係數從 CY2028＝FY2029 開始；HBM 價格 2027 三情境相同）"
PB = {}; r0 = 70
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028（FY2029）顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29（FY2029–30）顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "FY2030", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("dc", "  資料中心", "US$B", "#,##0.0"), ("cus", "    客製（XPU＋attach）", "US$B", "#,##0.0"),
                                ("ic", "    光互連", "US$B", "#,##0.0"), ("co", "  通訊與其他", "US$B", "#,##0.0"),
                                ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利（non-GAAP）", "US$B", "#,##0.0"),
                                ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2026: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2027 由下而上營收 − 錨點；FY2027 由下而上毛利 − 錨點（k_x、k_o、m_core 校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['bur']}-a27_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：毛利差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['gpo']}+D{PB['A']['mc']}*D{PB['A']['xc']}-a27_gp"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 70, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：Marvell 的 ASIC 每顆規格、光晶片組、校準值、份額加總 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — Marvell 的 XPU 與光晶片組規格（情境 A）。ASIC 核心售價、HBM 代購比例取自 Broadcom 模型（本專案權威，唯讀）；顆數：百萬顆（日曆年；與 HBM、台積電、網通、NVIDIA、Broadcom 模型同一份）；每顆價：US$K"
A_ = PB["A"]
CYS = (2026, 2027, 2028, 2029); cyc = lambda y: chr(ord("C") + y - 2026); fyof = lambda c: c + 1
hdr(UN, 3, ["項目", "單位", "CY2026（FY2027）", "CY2027（FY2028）", "CY2028（FY2029）", "CY2029（FY2030）", "說明"])
ur = 4
UN.cell(ur, 1, "1. XPU 每顆規格（情境 A；Marvell 營收口徑＝核心售價＋代購 HBM，乘 k_x）").font = BOLD; ur += 1
SPEC = {}
for k in XP:
    for key, lab, unit, f, fmt, note in (
        ("s", f"Marvell 份額：{PNAME[k]}", "%", lambda y: f"=PL!{FC[fyof(y)]}{A_['s_' + k]}", "0.0%", "In 頁 B 節（Assumed）"),
        ("b", f"Marvell 顆數：{PNAME[k]}", "百萬顆", lambda y: f"=PL!{FC[fyof(y)]}{A_['b_' + k]}", "0.000", "市場層顆數 × 份額"),
        ("h", f"每顆 HBM 金額：{PNAME[k]}", "US$K", lambda y: f"=PL!{FC[fyof(y)]}{A_['h_' + k]}", "0.00", "容量 × HBM 單價（情境 A）"),
        ("p", f"每顆 Marvell 營收（校準後）：{PNAME[k]}", "US$K", lambda y: f"=PL!{FC[fyof(y)]}{A_['a_' + k]}*PL!$D${A_['kx']}", "0.00", "（核心＋代購比例 × HBM）× k_x")):
        UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
        for y in CYS:
            UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
        UN.cell(ur, 7, note); SPEC[(key, k)] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "2. 份額加總檢查（Broadcom 模型份額＋Marvell 份額 ≤ 100%；其餘為 MediaTek、世芯、創意、華為等）").font = BOLD; ur += 1
for k in ("tpu", "oth"):
    UN.cell(ur, 1, f"Broadcom＋Marvell：{PNAME[k]}"); UN.cell(ur, 2, "%")
    for y in CYS:
        UN[f"{cyc(y)}{ur}"] = f"=av_{k}_{min(y, 2028) % 100}+PL!{FC[fyof(y)]}{A_['s_' + k]}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
    UN.cell(ur, 7, "Broadcom 份額讀 avgo 成品 In 頁"); SPEC[("sum", k)] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "3. 光晶片組（情境 A）").font = BOLD; ur += 1
for key, lab, unit, fmt in (("mod", "光模組數量", "百萬個", "0.0"), ("cs", "光晶片組市場", "US$B", "0.00"), ("oraw", "Marvell AI 光晶片組（未校準）", "US$B", "0.00"), ("ic", "Marvell 光互連（校準後，含相干）", "US$B", "0.00")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in CYS:
        UN[f"{cyc(y)}{ur}"] = f"=PL!{FC[fyof(y)]}{A_[key]}"; UN[f"{cyc(y)}{ur}"].number_format = fmt
    SPEC[("o", key)] = ur; ur += 1
UN.cell(ur, 1, "Marvell 光互連 ÷ 光晶片組市場"); UN.cell(ur, 2, "%")
for y in CYS:
    UN[f"{cyc(y)}{ur}"] = f"={cyc(y)}{SPEC[('o', 'ic')]}/{cyc(y)}{SPEC[('o', 'cs')]}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
SPEC[("o", "shr")] = ur; ur += 2
UN.cell(ur, 1, "4. 校準值（只用 FY2027 已發生的實績、當季指引與公司全年展望；不用共識、不用 FY2028 起的展望）").font = BOLD; ur += 1
CAL = {}
for key, lab, f, fmt, note in (("kx", "XPU 校準係數 k_x", f"=PL!D{A_['kx']}", "0.000", "FY2027 錨點 XPU ÷ 未校準由下而上；顆數、份額、核心單價的合併係數"),
                               ("ko", "光互連校準係數 k_o", f"=PL!D{A_['ko']}", "0.000", "（錨點光互連 − 相干）÷ 未校準 AI 光晶片組"),
                               ("mc", "XPU 核心毛利率 m_core", f"=PL!D{A_['mc']}", "0.0%", "讓 FY2027 毛利率等於錨點的殘差（Derived）"),
                               ("x27", "FY2027 錨點 XPU", "=a27_xpu", "0.00", "客製 − attach"),
                               ("i27", "FY2027 錨點光互連", "=a27_ic", "0.00", "資料中心差額")):
    UN.cell(ur, 1, lab); UN[f"C{ur}"] = f; UN[f"C{ur}"].number_format = fmt; UN.cell(ur, 7, note); CAL[key] = ur; ur += 1
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 17
UN.column_dimensions["G"].width = 70

# ---------------- Val：目標價、P/E ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元；non-GAAP EPS）。結論用「現價 ÷ FY2028 EPS（模型與共識）」與十年 P/E 規則值。GAAP EPS 多年為負（攤銷、併購），十年倍數改用「會計年度年底股價 ÷ 該年度 non-GAAP EPS」（Derived；與模型 EPS 同口徑）"
V["A2"] = "循環位置：共用規則（tools/cycle_position.py classify_lag，lag＝1：FY(t) 對照日曆年 t−1，與 NVIDIA 相同），見頂年度讀設備市場模型（" + "、".join(f"{s_} {PEAK[s_]}" for s_ in SC) + "）；FY2028 判定 " + "、".join(f"{s_} {CP.classify_lag(2028, PEAK[s_])}" for s_ in SC)
V.cell(4, 1, "0. 目標價與上檔空間").font = BOLD
hdr(V, 5, ["項目", "值", "", "", "說明"])
V.cell(6, 1, "目標價（12 個月，情境加權）")
V.cell(7, 1, "上檔空間（目標價 ÷ 現價 − 1）")
V.cell(8, 1, "情境 A 隱含價（樂觀）")
V.cell(9, 1, "情境 B 隱含價（基準）")
V.cell(10, 1, "情境 C 隱含價（下檔）")
V.cell(11, 1, "情境權重 A／B／C")
V.cell(12, 1, "目標價方法")
V["B6"] = "=tp_wA*B8+tp_wB*B9+tp_wC*B10"; V["B6"].number_format = "#,##0.0"
V["B7"] = "=B6/px-1"; V["B7"].number_format = "0.0%"
V["B11"] = "=tp_wA"; V["C11"] = "=tp_wB"; V["D11"] = "=tp_wC"
for c in "BCD": V[f"{c}11"].number_format = "0.00"
V["E11"] = "In 頁 tp_wA、tp_wB、tp_wC（Assumed；Andy 可改）"
V["E6"] = "＝ Σ 權重 × 第 3 節規則值隱含價"
V["B12"] = ("P/E 規則值：十年（CY2016–2025 年底＝FY2017–FY2026 會計年度年底）non-GAAP P/E 依 FY2028 循環位置取分位（中段＝中位數、見頂＝25 分位、谷底＝75 分位）"
            " × 各情境 FY2028 non-GAAP EPS（≈CY2027；隱含約 2028-01 的價格，未折現）。歷史倍數含 FY2020–FY2025 低 EPS 期的高倍數（偏高），也含 FY2018–FY2019、FY2023 的低倍數；照算不換倍數。"
            "情境標籤沿用工作單（A 樂觀、B 基準、C 下檔）；本模型 A＝市場層基準顆數，FY2028 三情境 EPS 相同，差異只來自 C 的循環位置（見頂 25 分位）")
V["B13"] = "=ROUND(tp_wA+tp_wB+tp_wC,6)"; V.cell(13, 1, "  檢查：權重合計（應為 1）")
for rr in range(6, 13):
    for c in range(1, 3): V.cell(rr, c).fill = KEY
V.cell(15, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 16, ["項目", "A", "B", "C", "說明"])
V.cell(17, 1, "現價（2026-10-23 收盤）"); V["B17"] = "=px"; V["B17"].number_format = "0.00"
V.cell(18, 1, "現價 ÷ FY2026 EPS（實績）")
V.cell(19, 1, "現價 ÷ FY2027 EPS（模型）")
V.cell(20, 1, "現價 ÷ FY2028 EPS（模型）")
V.cell(21, 1, "現價 ÷ FY2029 EPS（模型）")
V.cell(22, 1, "現價 ÷ FY2030 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((18, 2026), (19, 2027), (20, 2028), (21, 2029), (22, 2030)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(23, 1, "現價 ÷ 共識 FY2027 EPS"); V["B23"] = "=px/cons_eps27"; V["B23"].number_format = "0.0"
V.cell(24, 1, "現價 ÷ 共識 FY2028 EPS"); V["B24"] = "=px/cons_eps28"; V["B24"].number_format = "0.0"
V.cell(25, 1, "近 12 月 non-GAAP P/E（Q3 FY26～Q2 FY27 EPS）"); V["B25"] = "=px/(eps_q3_26+eps_q4_26+Q!C28+Q!D28)"; V["B25"].number_format = "0.0"
V.cell(26, 1, "近 12 月 GAAP P/E（Trefis，2026-10-16）"); V["B26"] = "=pe_ttm_gaap"; V["B26"].number_format = "0.0"
V["E19"] = "FY2027 為錨點（Q1、Q2 實績＋指引），三情境相同"; V["E20"] = "FY2028 三情境相同（顆數係數從 CY2028＝FY2029 開始）"
V.cell(28, 1, "2. 循環位置（共用規則 lag 版；FY(t) 對照 CY t−1）").font = BOLD
hdr(V, 29, ["情境", "設備市場見頂年度（CY）", "FY2028 判定", "FY2029 判定", "FY2030 判定", "規則"])
for i, s in enumerate(SC, 30):
    V.cell(i, 1, s); V.cell(i, 2, f"=pk_{s}")
    for j, y in enumerate((2028, 2029, 2030)):
        V.cell(i, 3 + j, "=" + CP.excel_rule_lag(str(y), f"pk_{s}"))
    V.cell(i, 6, "tools/cycle_position.py classify_lag（lag＝1）")
V.cell(34, 1, "3. 規則值隱含股價（美元；約 2028-01，未折現）＝ FY2028 EPS × 依 FY2028 循環位置的規則倍數").font = BOLD
hdr(V, 35, ["情境", "FY2028 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2028 EPS"])
IMP = {}
for j, s in enumerate(SC):
    i = 36 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f'=IF(C{30 + j}="見頂",pe_peak,IF(C{30 + j}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 4, f"=B{i}*C{i}").number_format = "#,##0"
    V.cell(i, 5, f"=D{i}/px-1").number_format = "0%"
    V.cell(i, 6, f"=px/B{i}").number_format = "0.0"
    for c in range(1, 7): V.cell(i, c).fill = KEY
    IMP[s] = i
V["B8"] = f"=D{IMP['A']}"; V["B9"] = f"=D{IMP['B']}"; V["B10"] = f"=D{IMP['C']}"
for rr in (8, 9, 10): V[f"B{rr}"].number_format = "#,##0.0"
V.cell(40, 1, "4. What must be true：2027-10 價格 ＝ 現價 ×（1＋報酬）所需的倍數（÷ FY2028 EPS；約一年後的價格看當時進行中的會計年度）").font = BOLD
hdr(V, 41, ["2027-10 價格 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2029 EPS，A）", "所需倍數（÷ FY2029 EPS，C）", "對照：十年 P/E 中位數"])
vr = 42
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!E{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；FY2029 欄反映 B、C 情境的顆數下修。")
WMT = vr
# 5. 十年倍數計算明細（輸入）
V.cell(50, 1, "5. 十年 P/E 計算明細：會計年度年底股價 ÷ 該年度 non-GAAP 稀釋 EPS（GAAP EPS 多年為負，不適用）").font = BOLD
hdr(V, 51, ["年底（CY）", "會計年度", "non-GAAP EPS", "1 月月底收盤（digrin，未調整）", "1 月 31 日股價（Macrotrends，含息調整）", "GAAP EPS（對照）", "來源"])
HIST = [("2016", "FY2017", 0.66, 14.87, 13.98, 0.04), ("2017", "FY2018", 1.19, 23.33, 22.24, 1.02), ("2018", "FY2019", 1.19, 18.53, 17.89, -0.30),
        ("2019", "FY2020", 0.66, 24.04, 23.45, 2.34), ("2020", "FY2021", 0.92, 51.46, 50.55, -0.41), ("2021", "FY2022", 1.57, 71.40, 70.41, -0.53),
        ("2022", "FY2023", 2.12, 43.15, 42.77, -0.19), ("2023", "FY2024", 1.51, 67.70, 67.42, -1.08), ("2024", "FY2025", 1.57, 112.86, 112.73, -1.02),
        ("2025", "FY2026", 2.84, 78.92, None, 3.07)]
HROW = {}
for i, (yy, fy, e, p1, p2, eg) in enumerate(HIST):
    rw = 52 + i; HROW[yy] = rw
    V.cell(rw, 1, yy); V.cell(rw, 2, fy)
    for col, v in ((3, e), (4, p1), (5, p2), (6, eg)):
        if v is not None:
            V.cell(rw, col, v).font = BLUE; V.cell(rw, col).number_format = "0.00"
    V.cell(rw, 7, "EPS：各年第四季 8-K 新聞稿（SEC EDGAR；FY2017–FY2019 為持續營業部門）" if i == 0 else ("Macrotrends 2026 年資料未找到" if yy == "2025" else ""))
V.cell(62, 1, "來源：EPS＝Marvell 8-K（FY2018 稿含 FY2017；FY2019 稿；FY2020 稿；FY2022 稿含 FY2021；FY2024 稿含 FY2023；FY2025、FY2026 為 Q4 FY26 稿）。股價：digrin.com（月資料，1 月；FY2018 年底實為 2018-02-03、FY2024 為 2024-02-03，取 1 月近似）；Macrotrends 1 月 31 日（含息調整，Interested-party 資料商）").font = Font(italic=True)
vals1 = [f"=D{{r}}/C{HROW[str(y)]}".replace("{{r}}", "{r}") for y in range(2016, 2026)]
vals1 = [f"=D{HROW[str(y)]}/C{HROW[str(y)]}" for y in range(2016, 2026)]
vals2 = [f"=E{HROW[str(y)]}/C{HROW[str(y)]}" if y < 2025 else None for y in range(2016, 2026)]
block_end = VR.add_rule_block(V, 65, [{"id": "P/E", "name": "P/E（non-GAAP）", "sources": [
    {"label": "digrin 1 月收盤 ÷ non-GAAP EPS", "vals": vals1, "tag": "Derived（股價：digrin，Interested-party 資料商；EPS：公司 Verified）"},
    {"label": "Macrotrends 1 月 31 日 ÷ non-GAAP EPS", "vals": vals2, "tag": "Derived（股價含息調整，略低；2025 未找到）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2028 EPS（A）", "basis": f"=PL!E{SUMR[('A', 'eps')]}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2028 EPS（A）", "basis": f"=PL!E{SUMR[('A', 'eps')]}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2028 EPS（A）", "basis": f"=PL!E{SUMR[('A', 'eps')]}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位。Marvell 會計年度 1 月底結束，CY(t) 年底對應 FY(t＋1) 年底（約一個月差）；GAAP P/E 多年為負（Macrotrends GAAP：2016 年 349.6x、2017 年 21.6x、2019 年 10.0x，其餘為負），改用 non-GAAP 自算兩序列（股價來源不同）。")
rule_rows = [rw for rw in range(65, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
V.column_dimensions["A"].width = 58
for c in "BCDEFG": V.column_dimensions[c].width = 18
V.column_dimensions["E"].width = 40; V.column_dimensions["G"].width = 60

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；k_x、k_o、m_core 在每個案例內依 FY2027 錨點重算，FY2027 恆等於錨點）：對 FY2028、FY2029 的營收、資料中心、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2028 營收", "FY2028 資料中心", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準", "FY2029 營收", "FY2029 資料中心", "FY2029 毛利率", "FY2029 EPS", "FY2029 EPS 相對基準"])
cases = [("XPU 顆數係數 ×0.8", "CY2027 起 Marvell XPU 顆數 ×0.8", dict(uf="uk_lo")), ("XPU 顆數係數 ×1.2", "×1.2", dict(uf="uk_hi")),
         ("其他 ASIC 份額 −3pt", "CY2027 起 Marvell 占其他 ASIC −3 個百分點", dict(sho="(-0.03)")), ("其他 ASIC 份額 +6pt", "+6 個百分點（Maia 放量較快）", dict(sho="0.06")),
         ("Trainium 份額 −5pt", "CY2027 起 −5 個百分點", dict(sht="(-0.05)")), ("TPU 份額 +5pt", "CY2027 起 +5 個百分點（Google 案提前）", dict(shp="0.05")),
         ("核心售價 ×0.8", "CY2027 起 XPU 核心售價 ×0.8", dict(pfc="pfc_lo")), ("核心售價 ×1.2", "×1.2", dict(pfc="pfc_hi")),
         ("HBM 代購比例 低", "代購比例取低值（k_x、m_core 重算）", dict(fh="f_hbm_lo")), ("HBM 代購比例 高", "代購比例取高值", dict(fh="f_hbm_hi")),
         ("光晶片組份額 −10pt", "CY2027 起 Marvell 占光晶片組 −10 個百分點", dict(shd="(-0.10)")), ("光晶片組份額 +5pt", "+5 個百分點", dict(shd="0.05")),
         ("每模組晶片組 ×0.8", "CY2027 起每模組光晶片組金額 ×0.8", dict(dspf="0.8")), ("線性直驅 +10pt", "CY2027 起 LPO 占比 +10 個百分點", dict(lpoa="0.10")),
         ("光互連 ×1.2", "CY2027 起 Marvell AI 光晶片組 ×1.2", dict(of="of_hi")),
         ("XPU attach FY2028 低", "FY2028 $0.6B", dict(a28="att28_lo")), ("XPU attach FY2028 高", "FY2028 $1.4B", dict(a28="att28_hi")),
         ("交換器 FY2028 年增 低", "+40%", dict(gsw28="g_sw28_lo")), ("通訊與其他 FY2028 −5%", "低值", dict(gco28="g_co28_lo")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2028, "#,##0.0"), ("D", "dc", 2028, "#,##0.0"), ("E", "gm", 2028, "0.0%"), ("F", "eps", 2028, "0.00"),
                           ("H", "rev", 2029, "#,##0.0"), ("I", "dc", 2029, "#,##0.0"), ("J", "gm", 2029, "0.0%"), ("K", "eps", 2029, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85；記憶體 B 路徑", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6；記憶體 C 路徑", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2028 三情境相同；顆數情境 B、C 的差異從 FY2029（CY2028）起。k_x、k_o 隨案例重算，所以 HBM 代購比例只改變營收組合與毛利率，不改 FY2027 水準。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司展望、投資人日目標、共識。只作事後對照，不為對上數字而調參數。公司 FY2028 起的展望（營收 $20B、資料中心 $18B、FY2029 客製 >$12B）為 Interested-party，不用於校準"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 營收 vs 公司展望（Interested-party；FY2028 起校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 公司展望 12", "=co_rev27", pl("rev", 2027), "公司（Q2 FY27 法說）", "錨點：第四季由全年展望推得，恆等")
rec("FY2028 營收 vs 投資人日 20", "=co_rev28", pl("rev", 2028), "投資人日（二手）", "公司展望一年內上修五次（$13B → $20B）")
rec("FY2028 資料中心 vs 投資人日 18", "=co_dc28", pl("dc", 2028), "Stocktwits 轉述（單一來源）", "")
rec("FY2028 客製 vs 「至少倍增」（FY2027 錨點 × 2）", "=2*a27_cus", pl("cus", 2028), "公司（Q2 FY27 法說）", "")
rec("FY2029 客製 vs 投資人日 >12", "=co_cus29", pl("cus", 2029), "投資人日（二手）", "含 Google 認股權證案的機率加權部分")
rec("FY2028 EPS vs 公司 3 月說法「遠高於 $5」", "=co_eps28", pl("eps", 2028), "Q4 FY26 法說（當時營收展望 $15B）", "", fmt="0.00")
RC.cell(rr_, 1, "B. What must be true：要達到 FY2028 營收 $20B，非通訊部分需要是模型資料中心的幾倍").font = BOLD; rr_ += 1
rec("FY2028：所需資料中心倍數（＝(20 − 模型通訊) ÷ 模型資料中心）", None, lambda s: f"=(co_rev28-PL!E{PB[s]['co']})/PL!E{PB[s]['dc']}", "—", "倍數 >1＝公司展望需要比市場層更多的 Marvell 出貨", fmt="0.00")
rec("FY2028：差額（US$B）", None, lambda s: f"=co_rev28-PL!E{PB[s]['rev']}", "—", "若全由客製補上，所需客製 ÷ 模型客製見下一列")
rec("FY2028：所需客製倍數（差額全由客製補）", None, lambda s: f"=(co_rev28-PL!E{PB[s]['rev']}+PL!E{PB[s]['cus']})/PL!E{PB[s]['cus']}", "—", "", fmt="0.00")
RC.cell(rr_, 1, "C. 共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "StockAnalysis", "錨點")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "StockAnalysis", "錨點（Q4 依全年展望與 In 頁假設）", fmt="0.00")
rec("FY2028 營收 vs 共識", "=cons_rev28", pl("rev", 2028), "StockAnalysis", "")
rec("FY2028 EPS vs 共識", "=cons_eps28", pl("eps", 2028), "StockAnalysis", "差距主要來自營收（共識已接近公司 $20B）", fmt="0.00")
rec("FY2028 共識隱含毛利率（EPS × 股數 ÷（1−稅）＋利息＋營業費用，÷ 共識營收）", "=(cons_eps28*PL!E" + str(PB['A']['sh']) + "/(1-tx_f)+int_f+PL!E" + str(PB['A']['opx']) + ")/cons_rev28", pl("gm", 2028), "推導（共識無毛利率）", "以模型營業費用估算", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "D. 產品線（FY2027 錨點拆分的一致性）").font = BOLD; rr_ += 1
rec("FY2027 光互連年增 vs 公司「超過 70%」", "=co_icg27", lambda s: "=Q!G14", "公司（Q1 FY27 法說，二手）", "光互連為資料中心差額；差額含 scale-up 光學、相干 DSP", fmt="0.0%", pt=True)
rec("FY2027 資料中心年增 vs 公司約 60%", "=co_dcg27", lambda s: "=Q!B35", "公司（Q2 FY27 法說）", "錨點恆等", fmt="0.0%", pt=True)
rec("CY2026 Marvell 光互連 ÷ 光晶片組市場", None, lambda s: f"=Units!C{SPEC[('o', 'shr')]}", "—", "含相干與 k_o；>100% 表示市場層光模組金額偏低或光互連含 AI 後端以外", fmt="0.0%")
rec("CY2026 光模組：本專案 vs LightCounting 26", "=rep_opt26", lambda s: "=plug_26", "markets/network；LightCounting（含 scale-up）", "")
rec("FY2028 客製 ÷ 資料中心", None, lambda s: f"=PL!E{PB[s]['cus']}/PL!E{PB[s]['dc']}", "—", "投資人日：FY2031 客製約 $30B ÷ 營收約 $80B", fmt="0.0%")
RC.cell(rr_, 1, "E. 校準").font = BOLD; rr_ += 1
rec("k_x（XPU 校準係數）", None, lambda s: f"=PL!D{PB[s]['kx']}", "FY2027 錨點 ÷ 未校準", "接近 1＝市場層顆數 × 份額 × Broadcom 模型單價與公司客製營收一致", fmt="0.000")
rec("k_o（光互連校準係數）", None, lambda s: f"=PL!D{PB[s]['ko']}", "（錨點光互連 − 相干）÷ 未校準", "高於 1：光互連包含 AI 後端可插拔以外的產品，或市場層光學偏低", fmt="0.000")
rec("m_core（XPU 核心毛利率，Derived）", None, lambda s: f"=PL!D{PB[s]['mc']}", "FY2027 毛利殘差", "客製毛利率低於公司平均（公司說法：客製稀釋毛利率；二手）", fmt="0.0%")
rec("FY2028 毛利率 vs FY2027 錨點（百分點）", "=a27_gp/a27_rev", pl("gm", 2028), "錨點", "公司：FY2028 與 FY2027 期末同區間（約 58%）", fmt="0.0%", pt=True)
rec("FY2030 毛利率 vs 公司 FY2031 目標 56–59% 中點", "=co_gm31", pl("gm", 2030), "投資人日（二手）", "模型偏低主要來自代購 HBM（Broadcom 模型的代購比例 0.6）；代購比例低時見 Sens", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "F. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 Marvell 客製營收 ÷ 部署 GW 對照 Tokenomics 非 NVIDIA 加速器比例"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 80
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 40; RC.column_dimensions["J"].width = 70

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝Marvell 8-K 新聞稿，已逐份抓取；法說與投資人日內容只經二手轉述）"; SO["A1"].font = BOLD
srcs = [("公司：Q4 與 FY2026 新聞稿（8-K）", Q4PR), ("公司：Q1 FY27 新聞稿（8-K）", Q1PR), ("公司：Q2 FY27 新聞稿（含 Q3 指引、六個月合計）", Q2PR),
        ("公司：Q3 FY26 新聞稿（舊五類終端市場拆分）", Q3PR26),
        ("公司：歷年第四季新聞稿（non-GAAP EPS：FY2017–FY2025）", EDG + "000183563224000004/q424_8kx232024ex-991.htm；" + EDG + "000183563222000006/q422_8kx1292022ex-991.htm；https://www.sec.gov/Archives/edgar/data/1058057/000105805720000007/q4208k212020ex-991.htm；https://investor.marvell.com/sec-filings/all-sec-filings/content/0001058057-19-000007/q419_8kx222019ex-991.htm；https://investor.marvell.com/sec-filings/all-sec-filings/content/0001058057-18-000003/q418_8kx232018ex-991.htm"),
        ("二手：Q2 FY27 法說逐字稿（FY2027 $12B、FY2028 $18B、資料中心 +60%、客製 FY2028 倍增）", "https://singjupost.com/transcript-marvell-tech-mrvl-q2-fy27-earnings-call-conference/"),
        ("二手：Q4 FY26 法說逐字稿（客製 $1.5B、交換器、AEC、XPU attach）", "https://www.roic.ai/quote/MRVL/transcripts/2026-year/4-quarter"),
        ("二手：Futurum Q1 FY27 摘要（光互連 >70%）", "https://futurumgroup.com/insights/marvell-q1-fy-2027-raises-full-year-outlook-on-ai-data-center-demand/"),
        ("二手：Zacks Q2 FY27 摘要", "https://www.zacks.com/stock/news/2981826/mrvl-q2-earnings-beat-on-data-center-strength-outlook-raised"),
        ("二手：投資人日（Futurum、Benzinga、Stocktwits）", "https://futurumgroup.com/?p=96506；https://www.benzinga.com/node/62198187；https://stocktwits.com/news-articles/markets/equity/marvell-investor-day-mrvl-stock-rockets-as-firm-lays-out-path-to-up-to-90-b-in-annual-revenue-by-fiscal-2031/cZDtJMyRBlE"),
        ("二手：Dolphin Research（Trainium 3 主供應商為世芯）", "https://longportapp.cn/topics/39096551"),
        ("二手：Motley Fool／Webull（Maia 由 Marvell 設計、Maia 300 多 GW）", "https://www.webull.com/news/15757166587814912"),
        ("二手：LightCounting PAM4 與相干 DSP 摘要（2026-04）", "https://lightcounting.com/newsletter/en/february-2026-pam4-and-coherent-dsps-updated-april-2026-381"),
        ("二手：StockAnalysis 共識與現價（2026-10-23）", "https://stockanalysis.com/stocks/mrvl/forecast/"),
        ("二手：Trefis（近 12 月 P/E、股價回落、下次財報約 2026-11-30）", "https://www.trefis.com/articles/619152/what-just-happened-to-marvell-stock/2026-10-19"),
        ("二手：digrin 月股價；Macrotrends P/E 與股價", "https://www.digrin.com/stocks/detail/MRVL/price；https://www.macrotrends.net/stocks/charts/MRVL/marvell-technology/pe-ratio"),
        ("市場層：markets/memory/market.py（UNITS）；tsmc、markets/network、avgo 成品（本 repo，唯讀）", ""),
        ("待補（人工）：客製各客戶營收；光互連實際金額；光 DSP 份額（LightCounting 付費）；HBM 代購比例；前瞻 P/E 十年歷史", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 80; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Marvell（MRVL）收支與估值模型 v0.1（AI 半導體 Project：運算層，客製 ASIC＋光互連 DSP＋交換器）　2026-10-25", True),
 ("版本：v0.1（初版）。樣板 avgo/；ASIC 每顆規格（核心售價、HBM 代購）讀 Broadcom 模型（本專案權威），本模型只加 Marvell 份額與光晶片組。本版未做先驗。", False),
 ("模型命題：Marvell 的資料中心營收＝客製 XPU（市場層 ASIC 顆數 × Marvell 份額 × 每顆售價）＋光互連（AI 光模組數 × 每模組晶片組金額 × 份額）＋交換器與其他；公司一年內把 FY2028 營收展望從 $13B 上修到 $20B，市場層顆數推得出多少？", False),
 ("核心問題：以市場層顆數與光模組市場往前推，Marvell FY2028 營收與 EPS 離公司 $20B 展望與共識有多遠；差距要靠哪條產品線補上，現價隱含多少。", False), ("", False),
 ("會計年度：FY(t)＝t−1 年 2 月～t 年 1 月底，FY(t) 直接取 CY(t−1)（與 NVIDIA 相同；不跨年加權）。FY2026A 實績；FY2027 錨點；預測 FY2028–FY2030（≈CY2027–29）。口徑：公司 non-GAAP（排除股票薪酬、攤銷、重組）。下次財報：Q3 FY27 約 2026-11-30～12 月初。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("客製 XPU：Trainium、其他 ASIC（Maia 等）、Google TPU 的市場層顆數（markets/memory × 情境係數）× Marvell 份額（Assumed）×（核心售價＋代購比例 × HBM 金額；Broadcom 模型規格）× k_x（FY2027 校準）＋ XPU attach。", False),
 ("光互連：markets/network 的 AI 後端可插拔光模組金額 ÷ 平均單價 ＝ 模組數 ×（1 − 線性直驅）× 每模組晶片組金額 × Marvell 份額 × k_o ＋ 相干 DSP／DCI。交換器、AEC／retimer、儲存：公司說法 × Assumed 成長；通訊與其他：FY2027 錨點 × 低個位數成長。", False),
 ("毛利率：光互連、其他資料中心、attach、通訊、代購 HBM 各用 Assumed 毛利率；XPU 核心毛利率 m_core 為 FY2027 殘差（Derived）。營業費用彈性 0.5（公司說法）、稅率 13%（公司說法）。", False),
 ("校準：FY2027 ＝ Q1、Q2 實績＋Q3 指引＋Q4（全年約 $12B − 前三季；資料中心年增約 60%）。不用共識、不用公司 FY2028 起的展望（只在 Recon 對帳）。", False),
 ("估值：目標價＝三情境規則值隱含價的加權（A 0.3、B 0.5、C 0.2；Andy 可改）；規則值＝十年 non-GAAP P/E 依 FY2028 循環位置取分位 × FY2028 EPS。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文＝SEC EDGAR 8-K；法說與投資人日內容只經二手轉述（標二手）。", False),
 ("已知限制：(1) k_x、k_o 是顆數、份額、單價的合併係數，無法分開識別；(2) Marvell 各平台份額、光晶片組份額與每模組金額皆為 Assumed；(3) 光互連 FY2027 錨點為資料中心差額，產品線拆分取公司口頭說法；(4) Google 認股權證案的平台未知（暫放 TPU）；(5) FY↔CY 近似；(6) 十年倍數含 FY2020–FY2025 低 EPS 期的高倍數；(7) R4 產業部署路徑未建（Recon 預留列）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("MRVL_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    inrow = {}
    for rw in IN.iter_rows(min_row=4):
        if rw[0].value and rw[3].value is None and rw[1].value is None: continue
        if rw[0].value: inrow[rw[0].value] = rw[0].row
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "SPEC": {f"{k[0]}|{k[1]}": v for k, v in SPEC.items()}, "CAL": CAL,
               "CHK": CHK_ROW, "SENSN": SENS_N, "IMP": IMP, "IN": inrow, "RC_END": rr_, "BLOCK_END": block_end, "RULE": rule_rows, "WMT": WMT},
              open(os.environ["MRVL_ROWMAP"], "w"), ensure_ascii=False, indent=1)
