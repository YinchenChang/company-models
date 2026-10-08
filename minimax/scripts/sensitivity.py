# 敏感度：逐項改寫 Inputs（build_xlsx.py --set），以 LibreOffice 重算後讀取關鍵輸出；結果寫入主檔 Sensitivity 頁（快照值）。
# 用法：python3 scripts/sensitivity.py 主檔.xlsx
import sys, os, subprocess, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from recalc import recalc

OUT = [('R09', 2026, '2026 營收 $M'), ('R09', 2030, '2030 營收 $M'), ('P07', 2030, '2030 毛利率'),
       ('P23', 2030, '2030 覆蓋率（每 MW 營收÷全成本）'), ('P18', 2030, '2030 經調整淨利 $M'),
       ('F11', 2030, '至 2030 累計需新融資 $M'), ('F19', 2026, '首次需新融資年'), ('V13', 2030, '折現價值÷市值')]
CASES = [
    ('基準', []),
    ('合約價 $8M/MW', ['I02=8']), ('合約價 $20M/MW', ['I02=20']),
    ('token 量成長低（0.8/0.5/0.3/0.2）', ['I09@2027=0.8', 'I09@2028=0.5', 'I09@2029=0.3', 'I09@2030=0.2']),
    ('token 量成長高（4.0/1.5/0.9/0.6）', ['I09@2027=4.0', 'I09@2028=1.5', 'I09@2029=0.9', 'I09@2030=0.6']),
    ('API 單價每年 −50%', ['I08=-0.5']), ('API 單價持平', ['I08=0']),
    ('2H26 企業端年化 350', ['I10=350']), ('2H26 企業端年化 750', ['I10=750']),
    ('訓練 MW 2027 起不再成長', ['I17=0']), ('訓練 MW 成長 ×0.5', ['I17=*0.5']), ('訓練 MW 成長 ×1.5', ['I17=*1.5']),
    ('模型效率年增 15%', ['I06@2027=2.3', 'I06@2028=2.645', 'I06@2029=3.042', 'I06@2030=3.498']),
    ('模型效率年增 50%', ['I06@2027=3.0', 'I06@2028=4.5', 'I06@2029=6.75', 'I06@2030=10.125']),
    ('2H26 推論效率改善 1.0', ['I29=1.0']), ('2H26 推論效率改善 2.0', ['I29=2.0']),
    ('消費端算力比 −0.15', ['I13=+-0.15']), ('消費端算力比 +0.15', ['I13=+0.15']),
    ('其他費用率 ×0.7', ['I20=*0.7']), ('其他費用率 ×1.3', ['I20=*1.3']),
    ('員工數 ×0.8', ['I18=*0.8']), ('員工數 ×1.3', ['I18=*1.3']),
    ('層級：低層占 30%', ['I04=0.3']), ('層級：低層占 90%', ['I04=0.9']),
    ('可轉債全數轉股', ['I23=1']),
]


def read(path):
    wb = load_workbook(path, data_only=True)
    loc = {}
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str): loc.setdefault(v, (ws, r))
    res = []
    for rid, y, _ in OUT:
        ws, r = loc[rid]
        res.append(ws.cell(r, 4 + y - 2025).value)
    return res


def main(master):
    tmp = os.path.join(ROOT, 'out', 'sens'); os.makedirs(tmp, exist_ok=True)
    rows = []
    for i, (name, sets) in enumerate(CASES):
        f = os.path.join(tmp, f'c{i:02d}.xlsx')
        args = ['python3', os.path.join(ROOT, 'build_xlsx.py'), f]
        for s in sets: args += ['--set', s]
        subprocess.run(args, check=True, capture_output=True)
        r = recalc(f)
        if r['status'] != 'success': raise SystemExit(f'{name}: {r}')
        rows.append((name, ' '.join(sets) or '—', read(f)))
        print(name, [round(x, 3) if isinstance(x, float) else x for x in rows[-1][2]], flush=True)
    wb = load_workbook(master)
    ws = wb['Sensitivity']
    ws['A1'] = '敏感度（快照值：每列為改寫該項 Inputs 後由 LibreOffice 重算的結果；改 Inputs 後此頁不會自動更新，需重跑 scripts/sensitivity.py）'
    ws['A1'].font = Font(bold=True)
    hdr = ['情境', '改寫'] + [o[2] for o in OUT]
    for c, h in enumerate(hdr, 1):
        x = ws.cell(3, c, h); x.font = Font(bold=True); x.fill = PatternFill('solid', fgColor='DDE6F0')
    for i, (name, s, vals) in enumerate(rows):
        ws.cell(4 + i, 1, name); ws.cell(4 + i, 2, s)
        for j, v in enumerate(vals):
            x = ws.cell(4 + i, 3 + j, v)
            x.number_format = '0.0%' if OUT[j][0] in ('P07',) else ('0.00"x"' if OUT[j][0] in ('P23', 'V13') else ('0' if OUT[j][0] == 'F19' else '#,##0'))
    ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 40
    for c in 'CDEFGHIJ': ws.column_dimensions[c].width = 16
    wb.save(master)
    json.dump([(n, s, v) for n, s, v in rows], open(os.path.join(ROOT, 'out', 'sensitivity.json'), 'w'), ensure_ascii=False, indent=1)


if __name__ == '__main__':
    main(sys.argv[1])
