# Tokenomics 官方快照：讀 Tokenomics checkout 的 model/CURRENT 具名範圍 → data/tokenomics_snapshot_<版本>.json。
# 唯讀（不改 Tokenomics）。用法：
#   產生：python3 scripts/tk_snapshot.py --tk-dir ../tk --names data/tokenomics_names.txt --out data/tokenomics_snapshot_v19.0.json
#   核對：python3 scripts/tk_snapshot.py --check --tk-dir ../tk --snapshot data/tokenomics_snapshot_v19.0.json   （DIFF／MISSING 只報 WARN，結束碼 0）
import argparse, datetime, hashlib, json, re, subprocess, sys
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci

REF = re.compile(r"^(?:'([^']+)'|([^!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$")


def read(tk_dir, names):
    tk_dir = Path(tk_dir)
    cur = (tk_dir / 'model' / 'CURRENT').read_text(encoding='utf-8').strip()
    path = tk_dir / 'model' / cur
    wb = openpyxl.load_workbook(path, data_only=True)
    defined = {k: v.attr_text for k, v in wb.defined_names.items()}
    out, missing = {}, []
    for n in names:
        if n not in defined:
            missing.append(n); continue
        m = REF.match(defined[n])
        ws = wb[m.group(1) or m.group(2)]
        c1, r1 = ci(m.group(3)), int(m.group(4))
        c2 = ci(m.group(5)) if m.group(5) else c1
        r2 = int(m.group(6)) if m.group(6) else r1
        vals = [ws.cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        e = dict(ref=defined[n], label=str(ws.cell(r1, 1).value).split('　')[0].strip(), unit=ws.cell(r1, 2).value, values=vals)
        if ws.title == 'Interface' and c2 != c1:   # 世代與成本情境表頭（Interface 第 4、5 列）
            e['generation'] = [ws.cell(4, k).value for k in range(c1, c2 + 1)]
            e['scenario'] = [ws.cell(5, k).value for k in range(c1, c2 + 1)]
        out[n] = e
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    return cur, sha, out, missing


def commit_of(tk_dir):
    return subprocess.run(['git', '-C', str(tk_dir), 'rev-parse', 'HEAD'], capture_output=True, text=True).stdout.strip()


def names_of(path):
    return [l.strip() for l in Path(path).read_text(encoding='utf-8').splitlines() if l.strip() and not l.startswith('#')]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--tk-dir', required=True)
    ap.add_argument('--names'); ap.add_argument('--out'); ap.add_argument('--snapshot')
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--read-date', default=datetime.date.today().isoformat())
    a = ap.parse_args()
    if a.check:
        snap = json.load(open(a.snapshot, encoding='utf-8'))
        cur, sha, vals, missing = read(a.tk_dir, list(snap['names']))
        print(f"快照 {snap['file']}（{snap['version']}，{snap['commit'][:7]}） vs Tokenomics CURRENT {cur}（{commit_of(a.tk_dir)[:7]}）")
        if sha != snap['file_sha256']:
            print(f"::warning::TK 現行檔與快照不同（SHA-256 {sha[:8]}… vs {snap['file_sha256'][:8]}…）")
        for n in missing:
            print(f'::warning::TK 快照 MISSING：{n}')
        d = [n for n, e in vals.items() if e['values'] != snap['names'][n]['values']]
        for n in d:
            print(f'::warning::TK 快照 DIFF：{n}')
        if a.names and set(names_of(a.names)) != set(snap['names']):
            print('::warning::tokenomics_names.txt 與快照名稱不一致')
        print(f'OK {len(vals) - len(d)}／DIFF {len(d)}／MISSING {len(missing)}')
        return
    cur, sha, vals, missing = read(a.tk_dir, names_of(a.names))
    if missing:
        sys.exit(f'Tokenomics 現行版缺名稱：{missing}')
    m = re.search(r'_v(\d+(?:\.\d+)?)\.xlsx$', cur)
    snap = dict(version=f'v{m.group(1)}' if m else '?', file=cur, commit=commit_of(a.tk_dir), file_sha256=sha,
                read_date=a.read_date, names=vals)
    Path(a.out).write_text(json.dumps(snap, ensure_ascii=False, indent=1, sort_keys=True) + '\n', encoding='utf-8')
    print(a.out, snap['version'], snap['commit'][:7], sha[:8], len(vals), '名')


if __name__ == '__main__':
    main()
