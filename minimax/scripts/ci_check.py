# CI／本機驗收：重算後的 xlsx 錯誤值 0、Checks 合計 0；快照與 tokenomics_names.txt 一致；TK_Link 標示的 SHA-256 與快照相同。
# 用法：python3 scripts/ci_check.py 已重算.xlsx   （任一項不符以代碼 1 結束）
import json, os, sys
from openpyxl import load_workbook
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
from tk_snapshot import names_of

x = sys.argv[1]
wb = load_workbook(x, data_only=True)
bad = []
ck = {r[0]: r[2] for r in wb['Checks'].iter_rows(values_only=True) if r[0]}
for k in ('CHK_Sum', 'CHK_Errors'):
    if ck.get(k) != 0: bad.append(f'{k}={ck.get(k)}（應為 0）')
cur = open(os.path.join(ROOT, 'data', 'tokenomics_snapshot_CURRENT'), encoding='utf-8').read().strip()
snap = json.load(open(os.path.join(ROOT, 'data', cur), encoding='utf-8'))
if set(names_of(os.path.join(ROOT, 'data', 'tokenomics_names.txt'))) != set(snap['names']):
    bad.append('快照名稱與 tokenomics_names.txt 不一致')
a2 = wb['TK_Link']['A2'].value or ''
if snap['file_sha256'] not in a2 or snap['commit'] not in a2:
    bad.append('TK_Link 標示的提交或 SHA-256 與快照不符')
print(f"快照 {snap['version']}（{snap['commit'][:7]}，{snap['file_sha256'][:8]}…）｜Checks 合計 {ck.get('CHK_Sum')}｜錯誤值 {ck.get('CHK_Errors')}")
if bad:
    print('::error::' + '；'.join(bad)); sys.exit(1)
print('OK')
