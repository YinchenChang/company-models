# 共用：讀已重算的 xlsx，取得計算頁（依 ID）、SRC_MM、Inputs、Sensitivity 的值與顯示格式（HTML 產生器與 check_html.py 共用，確保兩邊格式一致）
from openpyxl import load_workbook

CALC = ('Demand', 'Revenue', 'Compute', 'PnL', 'PxQ', 'Funding', 'Valuation')
Y = [2025, 2026, 2027, 2028, 2029, 2030]

FM = {
    'f0': lambda x: '—' if x is None else f'{x:,.0f}',
    'f1': lambda x: '—' if x is None else f'{x:,.1f}',
    'f2': lambda x: '—' if x is None else f'{x:,.2f}',
    'f3': lambda x: '—' if x is None else f'{x:,.3f}',
    'pc': lambda x: '—' if x is None else f'{x*100:.0f}%',
    'pc1': lambda x: '—' if x is None else f'{x*100:.1f}%',
    'x2': lambda x: '—' if x is None else f'{x:.2f}x',
    'yr': lambda x: '—' if x is None else ('不需要' if x == 9999 else f'{x:.0f}'),
    'txt': lambda x: '' if x is None else str(x),
}


def load(xl):
    wb = load_workbook(xl, data_only=True)
    D = {'wb': wb, 'V': {}, 'SRC': {}, 'INP': {}, 'LH': {}, 'TKV': {}, 'SENS': {}}
    for ws in wb.worksheets:
        if ws.title in CALC:
            for r in range(6, ws.max_row + 1):
                k = ws.cell(r, 1).value
                if isinstance(k, str) and len(k) <= 4:
                    D['V'][k] = [ws.cell(r, c).value for c in range(4, 10)]
                    D['V'][k + '_label'] = ws.cell(r, 2).value
    ws = wb['SRC_MM']
    for r in range(5, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k: D['SRC'][k] = ws.cell(r, 5).value
    ws = wb['Inputs']
    for r in range(6, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k:
            D['INP'][k] = [ws.cell(r, c).value for c in range(4, 10)]
            D['LH'][k] = {'lo': ws.cell(r, 10).value, 'hi': ws.cell(r, 11).value}
    ws = wb['TK_Link']
    for r in range(5, 17):
        k = ws.cell(r, 1).value
        if isinstance(k, str) and k.startswith('TK'): D['TKV'][k] = ws.cell(r, 6).value
    ws = wb['Sensitivity']
    for r in range(4, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k: D['SENS'][k] = [ws.cell(r, c).value for c in range(3, ws.max_column + 1)]
    return D


def val(D, key, y=None):
    """key：計算頁 ID（需年度）、S 開頭＝SRC_MM、I 開頭＝Inputs（純量取 D 欄；逐年取年度欄）、's:情境名'＝Sensitivity（y 為欄序）、'tk'＝TK_Link!A2。"""
    if key == 'tk': return D['wb']['TK_Link']['A2'].value
    if key == 'readme': return D['wb']['README']['A1'].value
    if '#' in key:
        b, w = key.split('#'); return D['LH'][b][w]
    if key in D['TKV']: return D['TKV'][key]
    if key.startswith('s:'): return D['SENS'][key[2:]][int(y)]
    if key in D['V']: return D['V'][key][Y.index(int(y))]
    if key in D['SRC']: return D['SRC'][key]
    if key in D['INP']:
        row = D['INP'][key]
        return row[Y.index(int(y))] if y else row[0]
    raise KeyError(key)
