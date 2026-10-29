"""MAG 財報萃取（API 版，2026-10-27 起）。

用途：10-Q／10-K 與財報新聞稿（8-K 附件 99.1）公布後，把 MAG 模型要更新的欄位萃取成候選值 Excel，
供建置代理填 company.json 前、Andy 與 chat 端核對時使用。**不改 company.json**；候選值不是模型數值。

流程：
1. 讀 mag/extract/requests/*.json（每檔一個請求，見 README）。主要模式：讀 Andy 下載到 Drive `MAG財報下載/` 的文件（drive_ids）或 repo 內檔案（files）；
   兩者皆無時才連 SEC EDGAR（GitHub 雲端 IP 目前被 SEC 擋，2026-10-27）。
2. 欄位清單直接從 <公司>/company.json 讀：latestQuarter、ytdActual、callFacts、quarterly.actuals[目標季]、
   以及 calendar_q.py → ROLL_FIELDS；說明取自 scripts/fields_doc.py。附「上次值」只作口徑參考。
3. 同一份文件由兩個模型各自萃取一次（Batch API，五折），逐欄比對；不一致者標「需確認」。
4. 回測模式（mode=backtest）：把萃取值與 company.json 現值比對，量化準確率。
5. 輸出 mag/extract/out/<請求名>.xlsx 與 .md 摘要；處理過的請求移到 mag/extract/requests/done/。

只用標準函式庫＋openpyxl。對 SEC 只帶 GitHub 免回覆地址（Andy 2026-10-27 同意），不用個人信箱。
"""
import glob
import html
import html.parser
import json
import os
import re
import shutil
import sys
import time
import urllib.error
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
MODELS = os.environ.get("EXTRACT_MODELS", "claude-opus-5-5,claude-sonnet-5-5").split(",")
# SEC 要求 User-Agent 附聯絡 email；Andy 2026-10-27 08:11 同意用 GitHub 免回覆地址（不用個人信箱）。
UA = "company-models research tool YinchenChang@users.noreply.github.com"
CIK = {"microsoft": 789019, "alphabet": 1652044, "amazon": 1018724}
MAX_DOC_CHARS = 900_000          # 單份文件上限（約 25–30 萬 token，在 1M context 內）
POLL_MAX_MIN = int(os.environ.get("POLL_MAX_MIN", "80"))
STALE_BATCH_H = float(os.environ.get("STALE_BATCH_H", "6"))   # 續接時批次超過此時數仍未完成，就取消改走一般 API


# ---------- 抓文件 ----------
def http_get(url, binary=False):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Encoding": "identity"})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                data = r.read()
                return data if binary else data.decode("utf-8", errors="replace")
        except urllib.error.HTTPError as e:
            if e.code in (429, 503) and attempt < 2:
                time.sleep(5 * (attempt + 1))
                continue
            body = e.read()[:400].decode("utf-8", errors="replace")
            raise RuntimeError(f"HTTP {e.code} {url}｜回應開頭：{body}") from None
        finally:
            time.sleep(0.2)   # SEC 公平存取：每秒少於 10 次


class _Text(html.parser.HTMLParser):
    """把 HTML 轉成文字；表格儲存格以 ' | ' 分隔、列尾換行，保留表格可讀性。"""
    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in ("br", "p", "div", "tr", "li", "h1", "h2", "h3", "h4"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip -= 1
        elif tag in ("td", "th"):
            self.out.append(" | ")
        elif tag in ("tr", "p", "div", "table"):
            self.out.append("\n")

    def handle_data(self, d):
        if not self.skip:
            self.out.append(d)


def html_to_text(s):
    p = _Text()
    p.feed(s)
    t = html.unescape("".join(p.out)).replace("\xa0", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"( \| )+(\n)", r"\2", t)
    t = re.sub(r"\n\s*\n+", "\n", t)
    return t


def find_filings(company, period_end):
    """回傳 [(標籤, 網址, 申報日)]：該期 10-Q／10-K 主文件，與期末後 0–45 天內 8-K 的附件 99.1。"""
    cik = CIK[company]
    sub = json.loads(http_get(f"https://data.sec.gov/submissions/CIK{cik:010d}.json"))
    rec = sub["filings"]["recent"]
    rows = [dict(zip(rec.keys(), vals)) for vals in zip(*rec.values())]
    docs = []
    for r in rows:
        if r["form"] in ("10-Q", "10-K") and r.get("reportDate") == period_end:
            acc = r["accessionNumber"].replace("-", "")
            docs.append((r["form"], f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc}/{r['primaryDocument']}", r["filingDate"]))
            break
    pe = time.strptime(period_end, "%Y-%m-%d")
    for r in rows:
        if r["form"] != "8-K" or "2.02" not in (r.get("items") or ""):
            continue
        days = (time.mktime(time.strptime(r["filingDate"], "%Y-%m-%d")) - time.mktime(pe)) / 86400
        if 0 <= days <= 45:
            acc = r["accessionNumber"].replace("-", "")
            idx = json.loads(http_get(f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc}/index.json"))
            for it in idx["directory"]["item"]:
                n = it["name"].lower()
                if re.search(r"ex-?99", n) and n.endswith((".htm", ".html")):
                    docs.append(("8-K EX-99", f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc}/{it['name']}", r["filingDate"]))
            break
    return docs


def read_local(path):
    """讀 repo 內的文件：PDF 用 pdftotext -layout（保留表格欄位對齊），HTML 轉文字，其餘當純文字。"""
    import subprocess
    full = path if os.path.isabs(path) else f"{ROOT}/{path}"
    low = full.lower()
    if low.endswith(".pdf"):
        return subprocess.run(["pdftotext", "-layout", full, "-"], capture_output=True, check=True).stdout.decode("utf-8", "replace")
    raw = open(full, "rb").read().decode("utf-8", "replace")
    return html_to_text(raw) if low.endswith((".htm", ".html")) else raw


def drive_download(file_id, dest):
    """下載 Drive 上「知道連結者可檢視」的檔案（只用於公開財報文件；需 Andy 先開放該資料夾的連結分享）。"""
    url = f"https://drive.usercontent.google.com/download?id={file_id}&export=download&confirm=t"
    req = urllib.request.Request(url, headers={"User-Agent": "company-models mag-extract"})
    with urllib.request.urlopen(req, timeout=300) as r:
        data = r.read()
    if data[:5] != b"%PDF-" and b"<html" in data[:2000].lower():
        raise RuntimeError(f"Drive 檔案 {file_id} 無法下載（可能未開放連結分享）")
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    open(dest, "wb").write(data)
    return dest


# ---------- 欄位清單 ----------
def field_list_from_file(path):
    """欄位清單檔（2026-10-29 起）：不依賴 company.json，供 v0.3 架構的 MAG、新建模型（例：SpaceX）使用。
    檔案為 JSON 陣列：[{"path": 欄位鍵, "desc": 說明（含單位）, "group": 分組（可省）, "prior": 上次值（可省）}]。"""
    rows = json.load(open(f"{ROOT}/{path}", encoding="utf-8"))
    fields = []
    for r in rows:
        if not r.get("path"):
            raise ValueError(f"欄位清單檔有缺 path 的列：{r}")
        fields.append({"path": r["path"], "group": r.get("group", "fields_file"),
                       "prior": r.get("prior"), "desc": r.get("desc", "")})
    if len({f["path"] for f in fields}) != len(fields):
        raise ValueError("欄位清單檔的 path 有重複")
    return None, fields


def field_list(company, target_q):
    co = json.load(open(f"{ROOT}/{company}/company.json", encoding="utf-8"))
    sys.path.insert(0, f"{ROOT}/{company}")
    sys.path.insert(0, f"{ROOT}/{company}/scripts")
    import calendar_q  # noqa: E402
    desc = {}
    try:
        import contextlib
        import io
        with contextlib.redirect_stdout(io.StringIO()):   # fields_doc 匯入時會印出 README 表格
            import fields_doc  # noqa: E402
        desc = {p: f"{d}（{u}）" for p, d, u, *_ in fields_doc.F}
        for sec, tbl in (("latestQuarter", getattr(fields_doc, "LQ", {})), ("callFacts", getattr(fields_doc, "CF", {}))):
            for k, v in tbl.items():
                desc[f"{sec}.{k}"] = f"{v[0]}（{v[1]}）" if isinstance(v, (tuple, list)) and len(v) >= 2 else str(v)
    except Exception as e:   # 說明檔讀不到時只少了說明，不影響萃取
        print(f"fields_doc 讀取失敗（只缺說明）：{e}")
    fields = []

    def leaf(section, obj):
        for k, v in obj.items():
            if k.startswith("_") or k in ("source", "date", "tag", "label", "throughQuarter", "months"):
                continue
            if isinstance(v, (int, float)) or v is None:
                p = f"{section}.{k}"
                fields.append({"path": p, "group": section.split(".")[0], "prior": v,
                               "desc": desc.get(p, desc.get(section, ""))})

    leaf("latestQuarter", co.get("latestQuarter", {}))
    leaf("ytdActual", co.get("ytdActual", {}))
    leaf("callFacts", co.get("callFacts", {}))
    for m in co.get("quarterly", {}).get("metrics", []):      # 季度層實際數的說明取自 metrics 定義
        desc[f"quarterly.actuals.{target_q}.{m['key']}"] = f"{m.get('actualLabel', m['label'])}（{m['unit']}）"
    if target_q and target_q in co.get("quarterly", {}).get("actuals", {}):
        leaf(f"quarterly.actuals.{target_q}", co["quarterly"]["actuals"][target_q])
    for cat, p, d in calendar_q.ROLL_FIELDS:
        try:
            prior = calendar_q._get_path(co, p)
        except Exception:
            continue                      # 本公司沒有此欄位
        fields.append({"path": p, "group": f"ROLL_FIELDS（{cat}）", "prior": prior, "desc": d})
    sys.path[:2] = []
    for m in ("calendar_q", "fields_doc"):
        sys.modules.pop(m, None)
    return co, fields


# ---------- 呼叫模型 ----------
SYSTEM = """你是財報資料萃取員。只根據提供的 SEC 文件原文萃取，不得推估、不得用記憶補值。
單位：金額一律 US$ 十億（bn），保留原文精度換算（例：$12,345 million → 12.345）；比例用小數（17.7% → 0.177）；股數用十億股；MW 用 MW。
若欄位清單附「上次值」，只用來對齊口徑與單位，不可抄寫。文件沒有直接揭露的欄位，value 填 null，並在 note 說明（例如「法說才有」「需由 X 與 Y 推算：X=…，Y=…」）。
ROLL_FIELDS 類欄位多為推導值或排程：只在文件直接揭露時填 value；否則 value 為 null，把可用的原始數字與出處寫在 note。
每筆附：doc（文件標籤）、loc（表名或段落標題，便於翻找）、quote（原文中含該數字的片段，15 字以內）。
只輸出一個 JSON 陣列，不要其他文字：[{"path":…,"value":數字或null,"unit":…,"doc":…,"loc":…,"quote":…,"note":…}]，path 必須逐一對應欄位清單，不可增減。"""


def build_prompt(req, docs_text, fields):
    hide = req.get("mode") == "backtest"     # 回測時上次值就是答案，必須隱藏，否則準確率會被抄寫灌水
    flist = "\n".join(f"- {f['path']}｜{f['group']}｜" + ("" if hide else f"上次值={json.dumps(f['prior'], ensure_ascii=False)[:200]}｜") + f"{f['desc'][:160]}"
                      for f in fields)
    return (f"公司：{req['company']}；目標期間：期末 {req['period_end']}（模型季度 {req.get('target_q')}）。\n"
            f"## 欄位清單（{len(fields)} 項）\n{flist}\n\n## 文件\n{docs_text}")


def api(path, body=None, method="GET"):
    req = urllib.request.Request(f"https://api.anthropic.com/v1/{path}",
                                 data=json.dumps(body).encode() if body is not None else None, method=method,
                                 headers={"x-api-key": API_KEY, "anthropic-version": "2023-06-01",
                                          "content-type": "application/json"})
    with urllib.request.urlopen(req, timeout=300) as r:
        return r.read().decode()


def run_batch(prompts, batch_id=None):
    """prompts：{custom_id: (model, prompt)} → ({custom_id: (文字, usage)}, 批次資訊)。
    batch_id 有值時只續接取回該批次（不重送）。逾時未完成時回傳空結果與批次資訊，由呼叫端保留請求以便下次續接。"""
    if batch_id:
        b = json.loads(api(f"messages/batches/{batch_id}"))
        age_h = (time.time() - time.mktime(time.strptime(b["created_at"][:19], "%Y-%m-%dT%H:%M:%S")) + time.timezone) / 3600
        if b["processing_status"] != "ended" and age_h > STALE_BATCH_H:
            # 批次異常緩慢：取消（未處理的請求不計費），改用一般 API 串流逐則送出（價格為 Batch 的 2 倍）
            print(f"批次 {batch_id} 已 {age_h:.1f} 小時未完成，取消並改用一般 API")
            api(f"messages/batches/{batch_id}/cancel", {}, "POST")
            out = {cid: stream_message(m, p) for cid, (m, p) in prompts.items()}
            return out, {"id": batch_id, "processing_status": "ended", "via": "standard"}
    else:
        reqs = [{"custom_id": cid, "params": {"model": m, "max_tokens": 32000, "system": SYSTEM,
                                             "messages": [{"role": "user", "content": p}]}}
                for cid, (m, p) in prompts.items()]
        b = json.loads(api("messages/batches", {"requests": reqs}, "POST"))
    print("batch", b["id"])
    t0 = time.time()
    while b["processing_status"] != "ended" and time.time() - t0 < POLL_MAX_MIN * 60:
        time.sleep(30)
        b = json.loads(api(f"messages/batches/{b['id']}"))
    out = {}
    if b["processing_status"] == "ended" and b.get("results_url"):
        for line in api(f"messages/batches/{b['id']}/results").splitlines():
            r = json.loads(line)
            res = r["result"]
            if res["type"] == "succeeded":
                msg = res["message"]
                out[r["custom_id"]] = ("".join(c.get("text", "") for c in msg["content"]), msg.get("usage", {}))
            else:
                out[r["custom_id"]] = (f"__error__ {json.dumps(res)[:500]}", {})
    return out, b


def stream_message(model, prompt):
    """一般 API（非 Batch）以串流取回一則回覆，回傳 (文字, usage)。長輸出用串流避免逾時。"""
    body = {"model": model, "max_tokens": 32000, "system": SYSTEM, "stream": True,
            "messages": [{"role": "user", "content": prompt}]}
    req = urllib.request.Request("https://api.anthropic.com/v1/messages", data=json.dumps(body).encode(),
                                 headers={"x-api-key": API_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"})
    text, usage = [], {}
    with urllib.request.urlopen(req, timeout=1800) as r:
        for raw in r:
            line = raw.decode("utf-8", "replace").strip()
            if not line.startswith("data:"):
                continue
            ev = json.loads(line[5:])
            t = ev.get("type")
            if t == "content_block_delta" and ev["delta"].get("type") == "text_delta":
                text.append(ev["delta"]["text"])
            elif t == "message_start":
                usage.update(ev["message"].get("usage", {}))
            elif t == "message_delta":
                usage.update(ev.get("usage", {}))
    return "".join(text), usage


def parse_json_array(text):
    m = re.search(r"\[.*\]", text, re.S)
    if not m:
        return []
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError:
        return []


# ---------- 比對與輸出 ----------
def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def close(a, b, tol=0.005):
    if num(a) and num(b):
        return abs(a - b) <= max(tol * max(abs(a), abs(b)), 0.0015)
    return a == b


def write_outputs(name, req, docs, fields, results, usage_total, batch_info):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    by_model = {m: {r.get("path"): r for r in parse_json_array(results.get(m, ("", {}))[0])} for m in MODELS}
    backtest = req.get("mode") == "backtest"
    wb = Workbook()
    ws = wb.active
    ws.title = "候選值"
    head = ["欄位路徑", "分組", "說明", "上次值（口徑參考）"]
    for m in MODELS:
        head += [f"{m} 值", f"{m} 位置", f"{m} 原文", f"{m} 備註"]
    head += ["兩模型一致", "判讀"] + (["回測：與現值一致"] if backtest else [])
    ws.append(head)
    stats = {"total": len(fields), "both": 0, "agree": 0, "disagree": 0, "one": 0, "none": 0, "bt_n": 0, "bt_ok": 0}
    bad, bt_grp = [], {}
    red, yellow = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="FFF3CD")
    for f in fields:
        vals = [by_model[m].get(f["path"], {}) for m in MODELS]
        v = [x.get("value") for x in vals]
        row = [f["path"], f["group"], f["desc"][:200], json.dumps(f["prior"], ensure_ascii=False)[:300]]
        for x in vals:
            row += [json.dumps(x.get("value"), ensure_ascii=False) if not num(x.get("value")) else x.get("value"),
                    f"{x.get('doc', '')}｜{x.get('loc', '')}", x.get("quote", ""), x.get("note", "")]
        given = [x for x in v if x is not None]
        if len(given) == len(MODELS):
            stats["both"] += 1
            ok = all(close(given[0], g) for g in given[1:])
            stats["agree" if ok else "disagree"] += 1
            agree, judge = ("是" if ok else "否"), ("" if ok else "需確認：兩模型數值不同")
        elif given:
            stats["one"] += 1
            agree, judge = "—", "需確認：只有一個模型找到"
        else:
            stats["none"] += 1
            agree, judge = "—", "兩模型皆未找到（見備註）"
        row += [agree, judge]
        if backtest:
            cand = given[0] if given and (len(given) == 1 or agree == "是") else None
            if num(f["prior"]) and cand is not None:
                stats["bt_n"] += 1
                hit = close(cand, f["prior"])
                stats["bt_ok"] += hit
                g = bt_grp.setdefault(f["group"], [0, 0])
                g[0] += 1
                g[1] += hit
                row.append("是" if hit else "否")
                if not hit:
                    bad.append((f["path"], f["prior"], cand))
            else:
                row.append("—")
        ws.append(row)
        if judge.startswith("需確認"):
            for c in ws[ws.max_row]:
                c.fill = red if "不同" in judge else yellow
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for col, w in zip("ABCD", (34, 16, 40, 18)):
        ws.column_dimensions[col].width = w

    s = wb.create_sheet("來源與用量")
    s.append(["文件", "網址", "申報日", "送出字數"])
    for lab, url, d, n in docs:
        s.append([lab, url, d, n])
    s.append([])
    s.append(["模型", "input tokens", "output tokens", "批次", "狀態"])
    for m in MODELS:
        u = results.get(m, ("", {}))[1]
        s.append([m, u.get("input_tokens"), u.get("output_tokens"), batch_info.get("id"), batch_info.get("processing_status")])
    std = batch_info.get("via") == "standard"
    s.append([("估計費用（一般 API；Opus 4/20、Sonnet 2/10 美元每百萬 token；批次逾時改走）" if std else
               "估計費用（Batch 五折；Opus 2/10、Sonnet 1/5 美元每百萬 token）"), round(usage_total, 3)])
    os.makedirs(f"{ROOT}/mag/extract/out", exist_ok=True)
    wb.save(f"{ROOT}/mag/extract/out/{name}.xlsx")

    md = [f"# 萃取摘要｜{name}", "",
          f"- 公司 {req['company']}；期末 {req['period_end']}；模式 {req.get('mode', 'live')}；模型 {'、'.join(MODELS)}",
          f"- 文件：" + "；".join(f"{lab}（{d}）" for lab, _, d, _ in docs),
          f"- 欄位 {stats['total']}：兩模型皆有值 {stats['both']}（一致 {stats['agree']}、不一致 {stats['disagree']}）；"
          f"只一個有值 {stats['one']}；皆無 {stats['none']}",
          f"- 估計費用約 US${usage_total:.2f}"]
    if backtest:
        rate = stats["bt_ok"] / stats["bt_n"] if stats["bt_n"] else 0
        md += [f"- **回測**：可比對 {stats['bt_n']} 項，與 company.json 現值一致 {stats['bt_ok']} 項（{rate:.0%}）"]
        md += [f"  - {g}：{ok}/{n}" for g, (n, ok) in bt_grp.items()]
        md += ["  - 註：ROLL_FIELDS 現值多為推導或排程，不一致不代表萃取錯；latestQuarter／ytdActual 才是直接可比的部分。", "",
               "| 欄位 | 現值 | 萃取值 |", "|---|---|---|"] + [f"| {p} | {a} | {b} |" for p, a, b in bad[:60]]
    md += ["", "候選值不是模型數值；填入 company.json 前須逐項對原文核對（見 xlsx「候選值」頁，紅＝兩模型不同、黃＝只一個找到）。"]
    open(f"{ROOT}/mag/extract/out/{name}.md", "w", encoding="utf-8").write("\n".join(md) + "\n")
    return stats


def main():
    if not API_KEY:
        print("未設定 ANTHROPIC_API_KEY，結束。")
        return
    reqs = sorted(glob.glob(f"{ROOT}/mag/extract/requests/*.json"))
    if not reqs:
        print("沒有待處理請求。")
        return
    for path in reqs:
        name = os.path.splitext(os.path.basename(path))[0]
        req = json.load(open(path, encoding="utf-8"))
        try:
            parts, meta = [], []
            local = list(req.get("files", []))
            for i, fid in enumerate(req.get("drive_ids", [])):
                local.append(drive_download(fid, f"{ROOT}/mag/extract/docs/_drive_{name}_{i}.pdf"))
            if local:   # 本地文件模式（2026-10-27 起的主要模式：SEC 擋雲端 IP，改由 Andy 手動下載）
                for doc in local:   # 不可用 path：外層迴圈的請求檔路徑（曾因覆蓋而把 PDF 搬成請求檔）
                    t = read_local(doc)[:int(req.get("max_doc_chars", MAX_DOC_CHARS))]
                    lab = os.path.basename(doc)
                    parts.append(f"\n\n===== 文件：{lab} =====\n{t}")
                    meta.append((lab, doc, req.get("filed", ""), len(t)))
            else:       # EDGAR 直連模式（雲端 IP 目前被 SEC 擋；保留供可直連的環境使用）
                docs = find_filings(req["company"], req["period_end"])
                if not docs:
                    raise RuntimeError("EDGAR 上找不到該期 10-Q／10-K 或財報 8-K（可能尚未申報）")
                for lab, url, d in docs:
                    t = html_to_text(http_get(url))[:MAX_DOC_CHARS]
                    parts.append(f"\n\n===== 文件：{lab}（申報日 {d}）=====\n{t}")
                    meta.append((lab, url, d, len(t)))
            if req.get("fields_file"):   # 欄位清單檔優先（v0.3 架構與新建模型沒有 company.json）
                co, fields = field_list_from_file(req["fields_file"])
            else:
                co, fields = field_list(req["company"], req.get("target_q"))
            prompt = build_prompt(req, "".join(parts), fields)
            results, binfo = run_batch({m: (m, prompt) for m in MODELS}, req.get("batch_id"))
            if binfo.get("processing_status") != "ended":     # 未完成：記下批次編號、保留請求，下次執行續接
                req["batch_id"] = binfo["id"]
                json.dump(req, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
                os.makedirs(f"{ROOT}/mag/extract/out", exist_ok=True)
                open(f"{ROOT}/mag/extract/out/{name}.md", "w", encoding="utf-8").write(
                    f"# 萃取處理中｜{name}\n\n批次 {binfo['id']} 尚未完成（{binfo.get('processing_status')}）；請求已記下批次編號，下次執行會續接取回，不重送。\n")
                print("批次未完成，保留請求", name, binfo["id"])
                continue
            os.makedirs(f"{ROOT}/mag/extract/out/raw", exist_ok=True)   # 原始輸出存檔（除錯用）
            for m in MODELS:
                open(f"{ROOT}/mag/extract/out/raw/{name}_{m}.txt", "w", encoding="utf-8").write(results.get(m, ("", {}))[0])
            price = {"opus": (2, 10), "sonnet": (1, 5), "haiku": (0.05, 0.25)}
            if binfo.get("via") == "standard":
                price = {k: (a * 2, b * 2) for k, (a, b) in price.items()}
            cost = 0.0
            for m in MODELS:
                u = results.get(m, ("", {}))[1]
                pi, po = next(v for k, v in price.items() if k in m)
                cost += (u.get("input_tokens", 0) * pi + u.get("output_tokens", 0) * po) / 1e6
            stats = write_outputs(name, req, meta, fields, results, cost, binfo)
            print(name, stats)
            os.makedirs(f"{ROOT}/mag/extract/requests/done", exist_ok=True)
            shutil.move(path, f"{ROOT}/mag/extract/requests/done/{name}.json")
        except Exception as e:
            os.makedirs(f"{ROOT}/mag/extract/out", exist_ok=True)
            open(f"{ROOT}/mag/extract/out/{name}.md", "w", encoding="utf-8").write(
                f"# 萃取失敗｜{name}\n\n{type(e).__name__}: {e}\n\n請求檔保留在 requests/，修正後重推即重跑。\n")
            print("失敗", name, e)


if __name__ == "__main__":
    main()
