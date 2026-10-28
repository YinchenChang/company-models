# SanDisk（SNDK US）收支與估值模型 v0.1（AI 半導體 Project：記憶體——NAND）
# 用法：python3 sndk/build_sndk.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度約 6 月底結束（52／53 週）：FY2027＝2026-07～2027-06 ≈ CY2026Q3–CY2027Q2。金額：十億美元（$B）；每股：美元。
# 結構比照 towa/build_towa.py（README、In、Q、Units、PL、Val、Recon、Sens、Sources）；估值主法比照記憶體樣板（tools/memory_model：Comparable 用前瞻 P/B）。
# 情境與行業一致（Andy 2026-10-25 13:13）：A／B／C 讀 markets/nand 同字母（價格＝markets/memory NAND_S、顆數＝SCEN_UNITS）；
# 公司特有變數（單位成本年降）A 最有利、C 最不利。份額由市場層提供（與 Kioxia 合資群份額一致），本檔不另訂。
import sys, os, glob, json
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as GL
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); FADE = PatternFill("solid", fgColor="E7E6E6")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
SCN = {"A": "順利", "B": "平平", "C": "不如預期"}
def defname(name, ref): wb.defined_names[name] = DefinedName(name, attr_text=ref)
rel = lambda f: os.path.relpath(f, ROOT)
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern)); assert len(fs) == 1, (pattern, fs); return fs[0]

# ================= 公司設定 =================
NAME, TICKER, VEN = "SanDisk", "SNDK", "SanDisk"
U = "$B"; PXU = "USD"; MF = "#,##0"; EF = "#,##0.00"
FY0 = 2026                                   # 錨點年（實績）
FYS = [2026, 2027, 2028, 2029]
def cyq(fy, k):                              # 會計年度 fy 第 k 季 → 日曆季
    return f"{fy - 1}Q{k + 2}" if k <= 2 else f"{fy}Q{k - 2}"
QL = [(fy, k) for fy in FYS for k in (1, 2, 3, 4)]   # 16 季：C..R
QT = {(2026, 1): "A", (2026, 2): "A", (2026, 3): "A", (2026, 4): "A"}   # 其餘＝D（驅動）
REFQ = (2027, 1)                             # 單位成本與營業費用的校準季（公司 FQ1 FY27 指引）

# ================= 上游（唯讀） =================
NAND_F = only("markets/nand/dist/*.xlsx"); MU_F = only("micron/dist/*.xlsx")
wn = openpyxl.load_workbook(NAND_F, data_only=True)
CYQS = [f"{y}Q{q}" for y in range(2025, 2031) for q in (1, 2, 3, 4)]
UP = {}
for r in wn["Mkt"].iter_rows():
    lab = str(r[0].value or "")
    for s in SC:
        if lab == f"{VEN} 營收（情境 {s}）": UP[("rev", s)] = ([r[2 + j].value for j in range(24)], r[0].row)
        if lab.startswith(f"{VEN} 位元（情境 {s}"): UP[("bit", s)] = ([r[2 + j].value for j in range(24)], r[0].row)
        if lab == f"Kioxia＋SanDisk（合資群）份額（情境 {s}）": UP[("jv", s)] = ([r[2 + j].value for j in range(24)], r[0].row)
        if lab == f"份額：{VEN}" and ("sh", s) not in UP and all(("rev", t) in UP for t in SC[:SC.index(s) + 1]) and ("rev", s) not in UP: pass
assert len(UP) == 9, sorted(UP)
JVSH = {}
for r in wn["Ann"].iter_rows():
    if r[0].value == "Kioxia＋SanDisk（合資群）份額":
        JVSH[SC[len(JVSH)]] = ([r[3 + j].value for j in range(6)], r[0].row)
assert len(JVSH) == 3
AIS = {}                                       # 情境 → ([AI 企業級 SSD 位元], [非 AI 企業級 SSD 位元], 列)：CY2025..CY2030
for r in wn["Ann"].iter_rows():
    if r[0].value == "  需求：AI 企業級 SSD": AIS.setdefault(SC[len(AIS)], [[r[3 + j].value for j in range(6)], None, r[0].row])
    if r[0].value == "  需求：非 AI 企業級 SSD": AIS[SC[len(AIS) - 1]][1] = [r[3 + j].value for j in range(6)]
assert len(AIS) == 3 and all(v[1] for v in AIS.values())
wm = openpyxl.load_workbook(MU_F, data_only=True)
MUV = wm["Val"]
def mu_hist(title):
    for i in range(1, MUV.max_row + 1):
        if MUV.cell(i, 1).value == title:
            y0 = i + 2; assert MUV.cell(y0, 1).value == "2016" and MUV.cell(y0 + 9, 1).value == "2025"
            return [MUV.cell(y0 + j, 2).value for j in range(10)], f"{rel(MU_F)}｜Val!B{y0}:B{y0 + 9}"
    raise KeyError(title)
MU_PB, MU_PB_SRC = mu_hist("P/B（主）：十年歷史倍數")
MU_PE, MU_PE_SRC = mu_hist("P/E（對照）：十年歷史倍數")
MU_PE = [v if isinstance(v, (int, float)) else None for v in MU_PE]
MU_NIM = None
for r in wm["In"].iter_rows(min_row=4):
    if r[0].value == "nim10": MU_NIM = (r[4].value, f"{rel(MU_F)}｜In!E{r[0].row}")
assert MU_NIM

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = f"In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（$B）；每股：美元；股數：百萬股。每個輸入附標記與來源；找不到者寫「未找到」"
IN["A2"] = "標記：Verified（公司文件、交易所資料）／Interested-party（寫明是誰、立場）／Analogy（寫明對照與差異）／Assumed（給低／高區間）／Derived（寫明換算）。SemiAnalysis 未使用"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
Q4 = "SanDisk FQ4 FY26 新聞稿（8-K，2026-08-05）"; Q2 = "SanDisk FQ2 FY26 新聞稿（8-K，2026-01-29）"
MS = "MarketScreener 預估頁（2026-10-25 擷取；頁面為 10-23 收盤；會計年度 6 月）"
sec("A. 季度實績與公司指引（$B；EPS 美元）")
QA = {1: (2.308, 0.298, None, 0.176, 0.112, 0.75), 2: (3.025, 0.509, None, 1.065, 0.803, 5.15), 3: (5.950, 0.784, 0.551, 4.111, 3.615, 23.03), 4: (8.965, 0.846, 0.545, 7.037, 6.903, 43.97)}
for k, (rv, gm, ox, op, ni, eps) in QA.items():
    src = Q2 if k <= 2 else Q4
    inp(f"q{k}_rev", f"FQ{k} FY26 營收", U, None, rv, None, "Verified", src, "FY26 合計 20.248（FQ4 新聞稿）")
    inp(f"q{k}_gm", f"FQ{k} FY26 GAAP 毛利率", "%", None, gm, None, "Verified", src, "")
    inp(f"q{k}_op", f"FQ{k} FY26 GAAP 營業利益", U, None, op, None, "Verified", src, "")
    inp(f"q{k}_ni", f"FQ{k} FY26 GAAP 淨利", U, None, ni, None, "Verified", src, "")
    inp(f"q{k}_eps", f"FQ{k} FY26 GAAP 稀釋 EPS", PXU, None, eps, None, "Verified", src, "")
K10Q = "SanDisk 10-Q（FQ2 FY26，2026-01-02 止；FQ3 FY26，2026-04-03 止）營收依終端市場表"; K10K = "SanDisk 10-K FY2026（2026-07-03 止）"
QSEG = {1: (0.269, 1.387, 0.652), 2: (0.440, 1.678, 0.907), 3: (1.467, 3.663, 0.820), 4: (2.977, 5.432, 0.556)}
for k, (dc_, ed_, cs_) in QSEG.items():
    stag = "Verified" if k in (2, 3) else "Derived"
    ssrc = K10Q if k in (2, 3) else (K10Q + "：六個月 − FQ2" if k == 1 else K10K + " 全年 − 10-Q 九個月")
    inp(f"q{k}_dc", f"FQ{k} FY26 營收：資料中心", U, None, dc_, None, stag, ssrc, "官方分部：Datacenter／Edge／Consumer（依終端市場）")
    inp(f"q{k}_ed", f"FQ{k} FY26 營收：邊緣", U, None, ed_, None, stag, ssrc, "Edge＝PC、手機、車用等 OEM 用 client SSD 與嵌入式")
    inp(f"q{k}_cs", f"FQ{k} FY26 營收：消費", U, None, cs_, None, stag, ssrc, "Consumer＝記憶卡、隨身碟、零售 SSD")
inp("a26_dc", "FY2026 營收：資料中心", U, None, 5.153, None, "Verified", K10K + "；" + Q4, "")
inp("a26_ed", "FY2026 營收：邊緣", U, None, 12.160, None, "Verified", K10K + "；" + Q4, "")
inp("a26_cs", "FY2026 營收：消費", U, None, 2.935, None, "Verified", K10K + "；" + Q4, "")
inp("a26_rev", "FY2026 營收（2025-07～2026-07-03）", U, None, 20.248, None, "Verified", Q4 + "：資料中心 5.153、邊緣 12.160、消費 2.935", "MarketScreener 同值")
inp("a26_ni", "FY2026 GAAP 淨利", U, None, 11.433, None, "Verified", Q4, "MarketScreener 同值")
inp("a26_eps", "FY2026 GAAP 稀釋 EPS", PXU, None, 73.76, None, "Verified", Q4 + "（non-GAAP 70.88）", "MarketScreener 73.76")
inp("a26_op", "FY2026 GAAP 營業利益", U, None, 12.389, None, "Verified", Q4, "")
inp("a26_gm", "FY2026 GAAP 毛利率", "%", None, 0.715, None, "Verified", Q4, "")
inp("a25_rev", "FY2025 營收", U, None, 7.355, None, "Verified", Q4 + "（前期欄）", "MarketScreener 7,355")
inp("a25_ni", "FY2025 GAAP 淨利", U, None, -1.641, None, "Verified", Q4 + "（含商譽減損）", "MarketScreener −1,641")
inp("a25_eps", "FY2025 GAAP 稀釋 EPS", PXU, None, -11.32, None, "Verified", Q4, "")
inp("g_rev", "公司指引：FQ1 FY27（2026-07～09）營收中值", U, 10.3, 10.55, 10.8, "Interested-party", Q4 + "：$10.3–10.8B", "公司指引（立場：公司）；市場層 markets/nand 以此推 SanDisk 2026Q3 份額，因此本模型 FQ1 FY27 營收＝指引（恆等式）")
inp("g_gm", "公司指引：FQ1 FY27 GAAP 毛利率中值", "%", 0.83, 0.84, 0.849, "Interested-party", Q4 + "：GAAP 83.0–84.9%（non-GAAP 83.0–85.0%）", "用來校準單位成本（FQ1 FY27 為基準季）")
inp("g_ox", "公司指引：FQ1 FY27 GAAP 營業費用中值", U, 0.574, 0.594, 0.614, "Interested-party", Q4 + "：GAAP $574–614M（non-GAAP $520–540M）", "")
inp("g_eps", "公司指引：FQ1 FY27 non-GAAP 稀釋 EPS 中值", PXU, 44, 45, 46, "Interested-party", Q4 + "：$44.00–46.00；稀釋股數約 155 百萬股、non-GAAP 稅率 15%", "")
inp("g_sh", "公司指引：FQ1 FY27 稀釋股數", "百萬股", None, 155, None, "Interested-party", Q4, "")

sec("B. 成本、費用、稅（公司特有；隨情境者 A 最有利、C 最不利）")
for s, (lo, b, hi), src in zip(SC, ((0.10, 0.14, 0.18), (0.08, 0.12, 0.16), (0.06, 0.10, 0.14)),
                               ("A：BiCS8→BiCS10（332 層，密度約 +59%）如期量產，每位元成本年降約 14%", "B：轉換略延後", "C：需求轉弱、稼動率下降，固定成本攤提使年降縮小")):
    inp(f"cd_{s}", f"情境 {s}：每位元成本年降", "%", lo, b, hi, "Assumed", src, "以 FQ1 FY27 指引毛利率校準的單位成本為起點；歷史每位元成本年降約 15–20%（Analogy：產業製程轉換）" if s == "A" else "同上")
inp("ox_var", "營業費用變動率（÷ 營收）", "%", 0.005, 0.01, 0.02, "Assumed", "FY26 GAAP 營業費用 2.083 ÷ 營收 20.248 ≈ 10%，但多為固定；FQ4 0.545 ÷ 8.965 ≈ 6%", "固定部分＝指引營業費用 − 變動率 × 指引營收（Derived）")
inp("ox_g", "營業費用固定部分年增", "%", 0.05, 0.10, 0.15, "Assumed", "研發（BiCS10、企業級 SSD 控制晶片）與獎金", "")
inp("tax", "有效稅率", "%", 0.12, 0.15, 0.20, "Interested-party", "公司指引 non-GAAP 稅率 15%（FQ1 FY27）", "GAAP 稅率未指引")
inp("rf", "淨現金收益率", "%", 0.03, 0.04, 0.05, "Assumed", "美元短期利率", "")
for y, (dlo, db, dhi), (clo, cb, chi) in ((2027, (0.33, 0.38, 0.45), (0.04, 0.06, 0.09)), (2028, (0.33, 0.42, 0.50), (0.04, 0.06, 0.09)), (2029, (0.33, 0.45, 0.55), (0.04, 0.06, 0.09))):
    inp(f"dcsh_{y % 100}", f"FY{y} 資料中心占營收", "%", dlo, db, dhi, "Assumed", "FQ1–FQ4 FY26 依序 11.7%、14.5%、24.7%、33.2%（10-Q、10-K）；企業級 SSD 位元占產業需求由 CY2025 30% 升到 CY2027 59%（markets/nand 情境 A）", "低＝停在 FQ4 FY26 水準；分部拆分不影響合併營收與 EPS")
    inp(f"cssh_{y % 100}", f"FY{y} 消費占營收", "%", clo, cb, chi, "Assumed", "FQ1–FQ4 FY26 依序 28.2%、30.0%、13.8%、6.2%（零售通路價格調整落後合約價）", "邊緣＝1 − 資料中心 − 消費（Derived）")
inp("sd_k", "公司營收 ÷ 市場層 SanDisk 營收（TrendForce 口徑）", "x", None, "=a26_rev/SUM(Units!E5:H5)", None, "Derived", "FY2026 營收 ÷ 市場層 2025Q3–2026Q2 SanDisk 營收（Units 頁）", "約 1.0：TrendForce 以公司財報為準")

sec("C. 資本、股東回饋（FY2026 年底＝2026-07-03）")
inp("eq0", "FY2026 年底股東權益", U, None, 15.736, None, "Verified", Q4, "")
inp("nc0", "FY2026 年底淨現金（現金 − 負債）", U, None, 4.762, None, "Verified", Q4 + "：現金 4.762、負債 0（已全數清償）", "")
inp("sh0", "流通股數（FY2026 年底後）", "百萬股", None, 146.42, None, "Verified", "stockanalysis 統計頁（2026-10-25）：146.42M", "MarketScreener FY2027 股數 146.419M；FQ4 已買回 $4.524B")
inp("dil", "稀釋股數 − 流通股數", "百萬股", None, "=g_sh-sh0", None, "Derived", "公司指引稀釋約 155 − 流通 146.42", "各年固定")
inp("da_y", "折舊攤銷（公司本身；年）", U, 0.12, 0.16, 0.25, "Verified／Assumed", Q4 + "：FY26 0.149", "合資廠折舊在晶圓成本（營業成本）內")
inp("cap_r", "資本支出 ÷ 營收（含合資群分攤；FY2027 起）", "%", 0.03, 0.06, 0.09, "Interested-party", "Motley Fool 2026-08-29：SanDisk FY27 毛資本支出約營收 6%（含與 Kioxia 合資的分攤）", "FY26 毛資本支出 0.562（其中合資分攤 0.519；SanDisk FQ4 FY26 簡報，markets/equipment In 頁）")
inp("payout", "買回 ÷ 自由現金流", "%", 0.3, 0.5, 0.8, "Assumed", Q4 + "：新增 $14B 授權（合計剩餘 $15.5B）；FQ4 買回 $4.524B", "無現金股利")
inp("px", "股價（2026-10-23 收盤）", PXU, None, 1238.03, None, "Verified", "stockanalysis 歷史股價（2026-10-23 收盤 1,238.03）", "MarketScreener 參考價 1,238.03（同日）；10-09 為 1,581.82")
inp("bb_px", "未來買回均價", PXU, None, "=px", None, "Assumed", "＝現價", "")

sec("D. 共識與對帳參考（只放 Recon，不作校準）")
inp("cons_rev27", "共識營收 FY2027", U, 43.4, 49.135, 53.2, "Interested-party", MS + "：49,135；stockanalysis（21 位）平均 49.1、區間 43.4–53.2", "")
inp("cons_rev28", "共識營收 FY2028", U, None, 58.436, None, "Interested-party", MS, "")
inp("cons_eps27", "共識 EPS FY2027", PXU, 186.97, 217.4, 238.36, "Interested-party", MS + "：217.4；stockanalysis（21 位）non-GAAP 214.82、區間 186.97–238.36", "")
inp("cons_eps28", "共識 EPS FY2028", PXU, None, 267.6, None, "Interested-party", MS, "")
inp("cons_eps29", "共識 EPS FY2029", PXU, None, 281.7, None, "Interested-party", MS, "")
inp("cons_bv27", "共識每股淨值 FY2027 年底（2027-06）", PXU, None, 313.8, None, "Interested-party", MS, "")
inp("cons_bv28", "共識每股淨值 FY2028 年底（2028-06）", PXU, None, 553.1, None, "Interested-party", MS, "")
inp("cons_bv29", "共識每股淨值 FY2029 年底（2029-06）", PXU, None, 650.6, None, "Interested-party", MS, "")
inp("cons_ni27", "共識淨利 FY2027", U, None, 33.336, None, "Interested-party", MS, "")
inp("cons_ni28", "共識淨利 FY2028", U, None, 38.384, None, "Interested-party", MS, "")
IN_ROW = {}; r = 4
def flush():
    global r
    for srow in rows:
        if srow[0] == "sec":
            IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
        key, item, unit, lo, base, hi, tag, s1, s2 = srow
        IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
        for j, v in ((4, lo), (5, base), (6, hi)):
            if v is not None:
                IN.cell(r, j, v)
                if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
        if key.startswith("tp_") or key in TPK:
            for j in range(1, 10): IN.cell(r, j).fill = KEY
        IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
        defname(key, f"In!$E${r}"); IN_ROW[key] = r
        if lo is not None: defname(key + "_lo", f"In!$D${r}")
        if hi is not None: defname(key + "_hi", f"In!$F${r}")
        r += 1
    rows.clear()
TPK = set()
flush()
# 目標價參數（黃底；Andy 可改）
sec("E. 目標價參數（黃底；Val 頁第 0 節 DCF＋Comparable 與收斂法用；Assumed 者 Andy 可改）")
TPK.update(["tp_wA", "tp_wB", "tp_wC", "tp_r", "tp_g", "tp_roe", "tp_w_dcf", "nim10", "cyc_mA", "cyc_mC", "fade_mA", "fade_mB", "fade_mC", "fade_gA", "fade_gB", "fade_gC", "cons_tp"])
for s_, w_ in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s_}", f"目標價情境權重 {s_}（A 順利／B 平平／C 不如預期）", "%", None, w_, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%")
inp("tp_r", "折現率（股東權益要求報酬；美、台、日、韓同一值）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%）", "記憶體循環股的盈餘波動大於一般大型股；依工作單不改值")
inp("tp_g", "長期成長率（永續；收斂法 B 的 FY2032 營收年增）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "")
inp("tp_roe", "ROE（近年平均；FY2025–FY2026，上市不足三年）", "%", None, "=AVERAGE(-1.641/9.216,11.433/15.736)", None, "Derived", Q4 + "：FY2025 淨利 −1.641 ÷ 權益 9.216、FY2026 11.433 ÷ 15.736", "2025-02 自 WDC 分拆上市，FY2024 以前為 WDC 快閃記憶體事業（無獨立權益）；只用兩年")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53：DCF／Comparable 各半）", "")
inp("nim10", "十年平均淨利率（Analogy：美光 FY2016–FY2025）", "%", None, MU_NIM[0], None, "Analogy", MU_NIM[1], "SanDisk 獨立上市不足十年（2025-02 分拆），以同為記憶體循環股的美光十年平均代替；差異：美光含 DRAM、HBM，NAND 純度低")
inp("cyc_mA", "循環中段係數：情境 A 淨利率收斂目標 ÷ 十年平均", "x", None, 1.2, None, "Assumed", "記憶體樣板預設 1.2（循環股收斂到循環中段，不延續高峰利潤率）", "Andy 可改")
inp("cyc_mC", "循環中段係數：情境 C 淨利率收斂目標 ÷ 十年平均", "x", None, 0.8, None, "Assumed", "記憶體樣板預設 0.8", "Andy 可改")
for s_, f_, t_ in (("A", "=nim10*cyc_mA", "十年平均 × cyc_mA"), ("B", "=nim10", "＝十年平均"), ("C", "=nim10*cyc_mC", "十年平均 × cyc_mC")):
    inp(f"fade_m{s_}", f"情境 {s_}：FY2032 淨利率收斂目標（循環中段：{t_}）", "%", None, f_, None, "Derived", "主串指示：循環產業 FY2030 後收斂用十年平均（記憶體樣板的循環中段）", "FY2030 起由 FY2029 淨利率線性收斂，FY2032 達到")
for s_, f_, t_ in (("A", "=tp_g+0.02", "tp_g＋2pt"), ("B", "=tp_g", "＝tp_g"), ("C", "=tp_g-0.02", "tp_g−2pt")):
    inp(f"fade_g{s_}", f"情境 {s_}：FY2032 營收年增收斂目標（{t_}）", "%", None, f_, None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("cons_tp", "市場共識目標價（賣方平均；低／高＝最低／最高）", PXU, 1000, 2144, 3600, "Interested-party", "stockanalysis（S&P Global，25 位，2026-10-18 更新）：平均 2,144、中位數 2,150、區間 1,000–3,600", "MarketScreener 25 位平均 2,173（10-09 收盤時）")
flush()
for col, w in zip("ABCDEFGHI", [12, 60, 10, 9, 11, 9, 18, 90, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# ---------------- Units：市場層（markets/nand 成品，唯讀） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = f"Units — 市場層 markets/nand 成品的 {VEN} 季度營收與位元（三情境；唯讀抄錄，藍字）。來源：{rel(NAND_F)}"
UN["A2"] = "C..Z＝CY2025Q1..CY2030Q4；營收為 TrendForce 口徑（$B），位元為季度出貨 × 營收份額（EB）。合資群份額＝Kioxia＋SanDisk（與 kioxia/ 模型讀同一列）"
hdr(UN, 3, ["項目", "單位"] + CYQS + ["來源"])
UR = {}; ur = 4
for kind, lab, unit in (("rev", "營收", "$B"), ("bit", "位元", "EB"), ("jv", "合資群份額", "%")):
    for s in SC:
        vals, srow = UP[(kind, s)]
        UN.cell(ur, 1, f"{VEN} {lab}（情境 {s}）" if kind != "jv" else f"Kioxia＋SanDisk 合資群份額（情境 {s}）"); UN.cell(ur, 2, unit)
        for j, v in enumerate(vals):
            c = UN.cell(ur, 3 + j, v); c.font = BLUE; c.number_format = "0.0%" if kind == "jv" else "#,##0.000"
        UN.cell(ur, 27, f"{rel(NAND_F)}｜Mkt!C{srow}:Z{srow}")
        UR[(kind, s)] = ur; ur += 1
UN.cell(ur + 1, 1, "AI 企業級 SSD 占企業級 SSD 位元需求（日曆年；markets/nand Ann 頁）").font = BOLD
hdr(UN, ur + 2, ["項目", "單位"] + [f"CY{y}" for y in range(2025, 2031)] + ["來源"])
AIR = {}
for i, s in enumerate(SC):
    a_, n_, srow = AIS[s]; rr_ = ur + 3 + i
    UN.cell(rr_, 1, f"AI ÷ 企業級 SSD（情境 {s}）"); UN.cell(rr_, 2, "%")
    for j in range(6):
        c = UN.cell(rr_, 3 + j, float(a_[j]) / (float(a_[j]) + float(n_[j]))); c.font = BLUE; c.number_format = "0.0%"
    UN.cell(rr_, 9, f"{rel(NAND_F)}｜Ann!D{srow}:I{srow} ÷（同列＋下一列）"); AIR[s] = rr_
AICOL = {y: GL(3 + y - 2025) for y in range(2025, 2031)}
assert UR[("rev", "A")] == 4      # sd_k 公式引用 Units!E5:H5 → 改為情境 A 營收列
UN.column_dimensions["A"].width = 36; UN.column_dimensions["AA"].width = 70
UCOL = {q: GL(3 + i) for i, q in enumerate(CYQS)}
IN[f"E{IN_ROW['sd_k']}"] = f"=a26_rev/SUM(Units!{UCOL['2025Q3']}{UR[('rev', 'A')]}:{UCOL['2026Q2']}{UR[('rev', 'A')]})"

# ---------------- Q：季度實績 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 季度實績（$B；EPS 美元）與公司指引。會計年度約 6 月底結束；FQ1＝7–9 月"
hdr(Q, 3, ["項目", "FQ1 FY26", "FQ2 FY26", "FQ3 FY26", "FQ4 FY26", "FY2026", "FQ1 FY27 指引", "標記", "來源"])
for i, (lab, key, fm) in enumerate((("營收", "rev", "#,##0.000"), ("GAAP 毛利率", "gm", "0.0%"), ("GAAP 營業利益", "op", "#,##0.000"), ("GAAP 淨利", "ni", "#,##0.000"), ("GAAP 稀釋 EPS", "eps", EF)), 4):
    Q.cell(i, 1, lab)
    for k in range(1, 5): Q.cell(i, 1 + k, f"=q{k}_{key}").number_format = fm
    Q.cell(i, 8, "Verified"); Q.cell(i, 9, Q2 + "；" + Q4)
Q["F4"] = "=SUM(B4:E4)"; Q["F6"] = "=SUM(B6:E6)"; Q["F7"] = "=SUM(B7:E7)"; Q["F8"] = "=a26_eps"
Q["F5"] = "=SUMPRODUCT(B4:E4,B5:E5)/F4"
for c, f in (("F4", "#,##0.000"), ("F5", "0.0%"), ("F6", "#,##0.000"), ("F7", "#,##0.000"), ("F8", EF)): Q[c].number_format = f
Q["G4"] = "=g_rev"; Q["G5"] = "=g_gm"; Q["G8"] = "=g_eps"; Q["G4"].number_format = "#,##0.000"; Q["G5"].number_format = "0.0%"; Q["G8"].number_format = EF
Q.cell(9, 1, "檢查：四季營收、淨利合計 − FY2026（應約為 0）"); Q["B9"] = "=F4-a26_rev"; Q["C9"] = "=F7-a26_ni"
Q.cell(10, 1, "季增：營收"); [Q.cell(10, 1 + k, f"={GL(1 + k)}4/{GL(k)}4-1").__setattr__("number_format", "0.0%") for k in range(2, 5)]
Q.cell(12, 1, "FQ1 FY27（2026-07～09）財報 2026-10-29 盤後公布（公司 2026-09-29 公告）；本版 FY2027 第一季＝市場層推導（營收等於公司指引中值，為恆等式）")
Q.column_dimensions["A"].width = 44
for c in "BCDEFG": Q.column_dimensions[c].width = 13
Q.column_dimensions["I"].width = 70

# ---------------- 引擎 ----------------
QC = {q: GL(3 + i) for i, q in enumerate(QL)}            # C..R
FYC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
DEFP = dict(cd=None, oxv="ox_var", oxg="ox_g", tax="tax", gm="g_gm", pay="payout", capr="cap_r", pxm="1", shm="1")
def engine(ws, r0, title, s, **over):
    p = dict(DEFP); p.update(over); cd = p["cd"] or f"cd_{s}"
    R = {}; rr = [r0]
    ws.cell(r0, 1, title).font = BOLD; rr[0] += 1
    def row(key, label, unit, fs, fmt="#,##0.000", note=""):
        r_ = rr[0]; ws.cell(r_, 1, label); ws.cell(r_, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r_}"] = f; ws[f"{col}{r_}"].number_format = fmt
        if note: ws.cell(r_, 20, note)
        R[key] = r_; rr[0] += 1
    hdr(ws, rr[0], ["A. 季度", "單位"] + [f"FQ{k} FY{fy % 100}（{cyq(fy, k)}）" for fy, k in QL] + ["", "說明"]); rr[0] += 1
    DQ = [q for q in QL if q not in QT]
    ref = QC[REFQ]
    nq = lambda q: QL.index(q) - QL.index(REFQ)
    uc = lambda q: UCOL[cyq(*q)]
    row("um", "市場層營收（TrendForce 口徑，同字母情境）", U, {QC[q]: f"=Units!{uc(q)}{UR[('rev', s)]}*{p['shm']}" + (f"*{p['pxm']}" if nq(q) > 0 else "") for q in QL}, note="markets/nand；Sens 可乘份額係數、價格係數（FQ2 FY27 起）")
    row("bit", "位元（市場層；同字母）", "EB", {QC[q]: f"=Units!{uc(q)}{UR[('bit', s)]}*{p['shm']}" for q in QL})
    row("rev", "營收", U, {**{QC[q]: f"=q{q[1]}_rev" for q in QT}, **{QC[q]: f"={QC[q]}{R['um']}*sd_k" for q in DQ}}, note="實績季＝公司財報；之後＝市場層 × 公司口徑係數")
    fyc = lambda q: q[0] % 100
    row("dc", "資料中心", U, {**{QC[q]: f"=q{q[1]}_dc" for q in QT}, **{QC[q]: f"={QC[q]}{R['rev']}*dcsh_{fyc(q)}" for q in DQ}}, note="官方分部；實績季＝10-Q／10-K；之後＝營收 × 資料中心占比（In）")
    row("ai", "  資料中心：AI 相關（＝資料中心 × AI 企業級 SSD 占比）", U, {QC[q]: f"={QC[q]}{R['dc']}*(Units!{AICOL[q[0] - 1]}{AIR[s]}+Units!{AICOL[q[0]]}{AIR[s]})/2" for q in QL}, note="命題子列（Derived）：會計年度 ≈ 前後兩個日曆年平均；假設 SanDisk 資料中心組合與產業企業級 SSD 相同")
    row("dcn", "  資料中心：非 AI", U, {QC[q]: f"={QC[q]}{R['dc']}-{QC[q]}{R['ai']}" for q in QL})
    row("cs", "消費", U, {**{QC[q]: f"=q{q[1]}_cs" for q in QT}, **{QC[q]: f"={QC[q]}{R['rev']}*cssh_{fyc(q)}" for q in DQ}}, note="官方分部")
    row("ed", "邊緣", U, {**{QC[q]: f"=q{q[1]}_ed" for q in QT}, **{QC[q]: f"={QC[q]}{R['rev']}-{QC[q]}{R['dc']}-{QC[q]}{R['cs']}" for q in DQ}}, note="官方分部；預測季＝營收 − 資料中心 − 消費")
    row("qchk", "檢查：分部合計 − 營收（季）", U, {QC[q]: f"={QC[q]}{R['dc']}+{QC[q]}{R['ed']}+{QC[q]}{R['cs']}-{QC[q]}{R['rev']}" for q in QL}, fmt="0.000")
    row("uc", "每 EB 營業成本（FQ1 FY27 以指引毛利率校準，之後每年 × (1 − 成本年降)）", "$B/EB", {**{QC[q]: f"={QC[q]}{rr[0] + 1}/{QC[q]}{R['bit']}" for q in QT},
        ref: f"=(1-{p['gm']})*g_rev*sd_k/sd_k/{ref}{R['bit']}" if False else f"=(1-{p['gm']})*{ref}{R['rev']}/{ref}{R['bit']}",
        **{QC[q]: f"=${ref}{rr[0]}*(1-{cd})^({nq(q)}/4)" for q in DQ if nq(q) > 0}}, fmt="0.0000")
    row("cogs", "營業成本", U, {**{QC[q]: f"=q{q[1]}_rev*(1-q{q[1]}_gm)" for q in QT}, **{QC[q]: f"={QC[q]}{R['bit']}*{QC[q]}{R['uc']}" for q in DQ}})
    row("gm", "毛利率", "%", {QC[q]: f"=1-{QC[q]}{R['cogs']}/{QC[q]}{R['rev']}" for q in QL}, fmt="0.0%")
    row("ox", "營業費用", U, {**{QC[q]: f"=q{q[1]}_rev-q{q[1]}_rev*(1-q{q[1]}_gm)-q{q[1]}_op" for q in QT},
                             **{QC[q]: f"=(g_ox-{p['oxv']}*g_rev)*(1+{p['oxg']})^({nq(q)}/4)+{p['oxv']}*{QC[q]}{R['rev']}" for q in DQ}}, note="固定部分以 FQ1 FY27 指引營業費用校準（Derived）")
    row("op", "營業利益", U, {QC[q]: f"={QC[q]}{R['rev']}-{QC[q]}{R['cogs']}-{QC[q]}{R['ox']}" for q in QL})
    yr_rows = rr[0]; rr[0] += 1
    row("ni", "淨利", U, {**{QC[q]: f"=q{q[1]}_ni" for q in QT}, **{QC[q]: f"=({QC[q]}{R['op']}+rf*${FYC[q[0] - 1]}{{NC}}/4)*(1-{p['tax']})" for q in DQ}}, note="（營業利益＋上一年底淨現金 × 收益率 ÷ 4）×（1 − 稅率）")
    row("epsq", "季 EPS（稀釋）", PXU, {QC[q]: f"={QC[q]}{R['ni']}/(sh0+dil)*1000" for q in DQ}, fmt=EF)
    rr[0] += 1
    hdr(ws, rr[0], ["B. 年度（會計年度）", "單位", "FY2025A", "FY2026A", "FY2027", "FY2028", "FY2029", "", "", "", "", "", "", "", "", "", "", "", "", "說明"]); rr[0] += 1
    qs = lambda fy, key: f"=SUM({QC[(fy, 1)]}{R[key]}:{QC[(fy, 4)]}{R[key]})"
    F3 = (2027, 2028, 2029)
    row("yrev", "營收", U, {"C": "=a25_rev", **{FYC[y]: qs(y, "rev") for y in FYS}})
    row("ydc", "營收：資料中心", U, {FYC[y]: qs(y, "dc") for y in FYS})
    row("yai", "  資料中心：AI 相關", U, {FYC[y]: qs(y, "ai") for y in FYS})
    row("yed", "營收：邊緣", U, {FYC[y]: qs(y, "ed") for y in FYS})
    row("ycs", "營收：消費", U, {FYC[y]: qs(y, "cs") for y in FYS})
    row("ychk", "檢查：分部合計 − 合併營收（年；FY2026 另比 10-K 分部）", U, {FYC[y]: f"={FYC[y]}{R['ydc']}+{FYC[y]}{R['yed']}+{FYC[y]}{R['ycs']}-{FYC[y]}{R['yrev']}" + ("+ABS(D{0}-a26_dc)+ABS(D{1}-a26_ed)+ABS(D{2}-a26_cs)".format(R['ydc'], R['yed'], R['ycs']) if y == 2026 else "") for y in FYS}, fmt="0.000")
    row("ycogs", "營業成本", U, {FYC[y]: qs(y, "cogs") for y in FYS})
    row("ygm", "毛利率", "%", {FYC[y]: f"=1-{FYC[y]}{R['ycogs']}/{FYC[y]}{R['yrev']}" for y in FYS}, fmt="0.0%")
    row("yox", "營業費用", U, {FYC[y]: qs(y, "ox") for y in FYS})
    row("yop", "營業利益", U, {FYC[y]: qs(y, "op") for y in FYS})
    row("yom", "營業利益率", "%", {FYC[y]: f"={FYC[y]}{R['yop']}/{FYC[y]}{R['yrev']}" for y in FYS}, fmt="0.0%")
    row("yni", "淨利", U, {"C": "=a25_ni", **{FYC[y]: qs(y, "ni") for y in FYS}})
    row("da", "折舊攤銷（公司本身）", U, {FYC[y]: "=da_y" for y in FYS})
    row("cap", "資本支出（含合資群分攤）", U, {"D": 0.562, **{FYC[y]: f"={p['capr']}*{FYC[y]}{R['yrev']}" for y in F3}}, note="FY2026 毛資本支出 0.562（公司簡報；markets/equipment In 頁）")
    row("fcf", "自由現金流（淨利＋折舊 − 資本支出）", U, {FYC[y]: f"={FYC[y]}{R['yni']}+{FYC[y]}{R['da']}-{FYC[y]}{R['cap']}" for y in FYS})
    row("bb", "買回", U, {"D": 4.524, **{FYC[y]: f"=MAX(0,{p['pay']}*{FYC[y]}{R['fcf']})" for y in F3}}, note="FY2026 實績 4.524（FQ4）")
    row("eq", "期末權益", U, {"D": "=eq0", **{FYC[y]: f"={FYC[y - 1]}{rr[0]}+{FYC[y]}{R['yni']}-{FYC[y]}{R['bb']}" for y in F3}})
    row("nc", "期末淨現金", U, {"D": "=nc0", **{FYC[y]: f"={FYC[y - 1]}{rr[0]}+{FYC[y]}{R['fcf']}-{FYC[y]}{R['bb']}" for y in F3}})
    row("sh", "期末流通股數", "百萬股", {"D": "=sh0", **{FYC[y]: f"={FYC[y - 1]}{rr[0]}-{FYC[y]}{R['bb']}/bb_px*1000" for y in F3}}, fmt="#,##0.0")
    row("eps", "EPS（稀釋：期末流通股數＋dil）", PXU, {"C": "=a25_eps", "D": "=a26_eps", **{FYC[y]: f"={FYC[y]}{R['yni']}/({FYC[y]}{R['sh']}+dil)*1000" for y in F3}}, fmt=EF)
    row("bv", "每股淨值 BVPS（期末流通股數）", PXU, {FYC[y]: f"={FYC[y]}{R['eq']}/{FYC[y]}{R['sh']}*1000" for y in FYS}, fmt=EF)
    row("roe", "ROE（期末權益）", "%", {FYC[y]: f"={FYC[y]}{R['yni']}/{FYC[y]}{R['eq']}" for y in FYS}, fmt="0%")
    row("cfps", "每股自由現金流（稀釋）", PXU, {FYC[y]: f"={FYC[y]}{R['fcf']}/({FYC[y]}{R['sh']}+dil)*1000" for y in F3}, fmt=EF)
    row("chk", "檢查：FY2026 營收、淨利 − 實績（應為 0）", U, {"D": f"=ABS(D{R['yrev']}-a26_rev)+ABS(D{R['yni']}-a26_ni)"}, fmt="0.000")
    # 回填淨利公式中的上一年淨現金列
    for q in DQ:
        c = ws[f"{QC[q]}{R['ni']}"]; c.value = c.value.replace("{NC}", str(R["nc"]))
    R["end"] = rr[0]
    return R

PL = wb.create_sheet("PL")
PL["A1"] = f"PL — {NAME} 損益（$B；FY2026 實績；FY2027–FY2029 由市場層 markets/nand 驅動；FY2030–FY2032 收斂法見底部）。頂端摘要，下方三情境各一個完整推導區塊"
TITLES = {s: f"情境 {s}（{SCN[s]}）：markets/nand 情境 {s}（價格路徑 {s}、2028–29 顆數係數 {s}、合資群份額 {s}）；每位元成本年降 cd_{s}" for s in SC}
PB = {}; r0 = 56
for s in SC:
    PB[s] = engine(PL, r0, TITLES[s], s); r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026A", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}（{SCN[s]}）").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("yrev", "營收", U, "#,##0.00"), ("ydc", "  資料中心", U, "#,##0.00"), ("yai", "    資料中心：AI 相關", U, "#,##0.00"), ("yed", "  邊緣", U, "#,##0.00"), ("ycs", "  消費", U, "#,##0.00"), ("ygm", "毛利率", "%", "0.0%"), ("yop", "營業利益", U, "#,##0.00"), ("yom", "營業利益率", "%", "0.0%"),
                                ("yni", "淨利", U, "#,##0.00"), ("eps", "EPS", PXU, EF), ("bv", "每股淨值", PXU, EF), ("fcf", "自由現金流", U, "#,##0.00")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y, c in FYC.items():
            if PL[f"{c}{PB[s][key]}"].value is None: continue
            PL[f"{c}{sr}"] = f"={c}{PB[s][key]}"; PL[f"{c}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對：三情境 FY2026 − 實績（應為 0）"); PL[f"D{sr}"] = "=" + "+".join(f"D{PB[s]['chk']}" for s in SC); PL[f"D{sr}"].number_format = "0.000"; PL_CHK = sr
assert sr < 54
PL.column_dimensions["A"].width = 58; PL.column_dimensions["B"].width = 10; PL.column_dimensions["T"].width = 60
for c in "CDEFGHIJKLMNOPQR": PL.column_dimensions[c].width = 12

# --- 延伸到 FY2032（收斂法） ---
p0 = PL.max_row + 3
PL.cell(p0, 1, "延伸到 FY2032：FY2026–FY2029 引用本頁頂端摘要；灰底 FY2030–FY2032 為收斂法，非驅動因子推導").font = BOLD
PL.cell(p0 + 1, 1, "收斂法：營收年增由 FY2029 年增線性收斂到 FY2032 的長期成長（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt）；淨利率由 FY2029 線性收斂到循環中段（A＝十年平均 × 1.2、B＝十年平均、C＝十年平均 × 0.8；十年平均為美光 Analogy）；股數維持 FY2029")
EXTC = {y: GL(3 + y - 2026) for y in range(2026, 2033)}   # C＝FY2026 … I＝FY2032
hdr(PL, p0 + 2, ["情境（延伸）", "單位", "FY2026A", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
EXT = {}; pr = p0 + 3
for s in SC:
    PL.cell(pr, 1, f"情境 {s}（{SCN[s]}）").font = BOLD; pr += 1
    R = {k: pr + i for i, k in enumerate(("rev", "g", "nim", "ni", "eps"))}
    for k, lab in (("rev", "營收"), ("g", "  營收年增"), ("nim", "  淨利率（淨利 ÷ 營收）"), ("ni", "淨利（收斂法延伸）"), ("eps", "EPS（延伸）")):
        PL.cell(R[k], 1, lab); PL.cell(R[k], 2, {"rev": U, "ni": U, "eps": PXU}.get(k, "%"))
    for y in range(2026, 2033):
        c = EXTC[y]; pc = EXTC.get(y - 1)
        if y <= 2029:
            for k, kk in (("rev", "yrev"), ("ni", "yni"), ("eps", "eps")): PL[f"{c}{R[k]}"] = f"={FYC[y]}{SUMR[(s, kk)]}"
            PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
            if y > 2026: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"=$F{R['g']}+(fade_g{s}-$F{R['g']})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"=$F{R['nim']}+(fade_m{s}-$F{R['nim']})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['eps']}"] = f"=$F{R['eps']}*{c}{R['ni']}/$F{R['ni']}"
            for kk in R: PL[f"{c}{R[kk]}"].fill = FADE
        for kk, fm in (("rev", "#,##0.00"), ("g", "0.0%"), ("nim", "0.0%"), ("ni", "#,##0.00"), ("eps", EF)):
            if PL[f"{c}{R[kk]}"].value is not None: PL[f"{c}{R[kk]}"].number_format = fm
    PL.cell(R["rev"], 10, "FY2030 起為收斂法，非驅動因子推導")
    EXT[s] = R; pr += 6

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — 估值。第 0 節為目標價（DCF 與 Comparable 加權；Comparable 依記憶體慣例用前瞻 P/B）；第 5–6 節為目前倍數、十年倍數與規則值（舊法參考）"
V["A2"] = "SanDisk 2025-02 自 WDC 分拆上市，無十年自身倍數歷史：十年 P/B、P/E 以美光（同為記憶體循環股）為 Analogy，見第 6 節"
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(rr, key, label, fs, fmt, note=""):
    V.cell(rr, 1, label); T[key] = rr
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{rr}"] = f; V[f"{c}{rr}"].number_format = fmt
    if note: V[f"E{rr}"] = note
t0 = 4
V.cell(t0, 1, "0. 目標價").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD; T["h01"] = t0 + 1
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["hd1"] = t0 + 2
V.cell(t0 + 3, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{t0 + 3}"] = "PL 頁延伸區塊；FY2030 起為收斂法（灰底）"; T["eps_h"] = t0 + 3
EPSR = {}
for i, y in enumerate(YRS6):
    tset(t0 + 4 + i, f"eps{y}", f"  FY{y}", [f"=PL!{EXTC[y]}{EXT[s]['eps']}" for s in SC], EF); EPSR[y] = t0 + 4 + i
rr = t0 + 10
V.cell(rr, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = rr
V[f"E{rr}"] = "記憶體資本密集：FY2027–FY2029 改用每股（淨利＋折舊－資本支出）＝ PL 區塊每股自由現金流（資本支出含合資群分攤）；收斂期（FY2030 起）用 EPS ×（1 − 再投資率）"
tset(rr + 1, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8；FY2030 起用）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為 FY2025–FY2026 平均（In 頁 tp_roe）")
RV = rr + 1; CFR = {}
for i, y in enumerate(YRS6):
    fs = [(f"=PL!{FYC[y]}{PB[s]['cfps']}" if y <= 2029 else f"={c}{EPSR[y]}*(1-$B${RV})") for s, c in zip(SC, "BCD")]
    tset(rr + 2 + i, f"cf{y}", f"  FY{y}" + ("（淨利＋折舊－資本支出，每股）" if y <= 2029 else "（EPS ×（1 − 再投資率））"), fs, EF); CFR[y] = rr + 2 + i
c1, c6 = CFR[2027], CFR[2032]
tset(rr + 8, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], MF,
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027（FY2027 於 2027-06 結束，t 以會計年度計，略低估）；負值取 0（股東有限責任）")
tset(rr + 9, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{rr + 8}*(1+tp_r)" for c in "BCD"], MF)
tset(rr + 10, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{rr + 9}+tp_wB*C{rr + 9}+tp_wC*D{rr + 9}", None, None], MF, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁 tp_wA／tp_wB／tp_wC")
q = rr + 12
V.cell(q, 1, "0.2 Comparable（市場主流同業前瞻 P/B；記憶體慣例）").font = BOLD; T["h02"] = q
hdr(V, q + 1, ["同業與前瞻 P/B", "代號", "前瞻 P/B", "日期", "來源", "標記"]); T["phd"] = q + 1
PEERS = [("美光", "MU", 3.35, "2026-10-23", "MarketScreener：現價 ÷ FY2027E（2027-09）每股淨值（micron/ 模型同值）"),
         ("SK 海力士", "000660.KS", 1.35, "2026-10-23", "MarketScreener：現價 ÷ 2027E 每股淨值（skhynix/ 模型同值）"),
         ("三星電子", "005930.KS", 1.26, "2026-10-23", "MarketScreener：現價 ÷ 2027E 每股淨值（samsung/ 模型同值）"),
         ("Kioxia", "285A.T", "=12755/7808", "2026-10-23", "現價 ¥12,755（stockanalysis，分割後）÷ MarketScreener FY2028E（2028-03）每股淨值 ¥7,808")]
for i, (nm, tk, pb, dt, src) in enumerate(PEERS, q + 2):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pb); c.number_format = "0.00"
    if not isinstance(pb, str): c.font = BLUE
    V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
T["p1"] = q + 2; T["pl"] = q + 1 + len(PEERS)
q = T["pl"] + 1
tset(q, "pmed", "同業前瞻 P/B 中位數", [f"=MEDIAN(C{T['p1']}:C{T['pl']})", None, None], "0.00",
     "記憶體同業（美光、SK 海力士、三星、Kioxia；無 A 股）。WDC（硬碟）不列入 P/B：僅有落後 P/B 13.5x（stockanalysis），與記憶體不可比；P/E 並列參考見下方。SanDisk 自身現價 ÷ 2027 年底每股淨值（FY27、FY28 共識平均）約 2.9x")
tset(q + 1, "fbv", "12 個月後的前瞻每股淨值（情境加權）", ["=" + "+".join(f"tp_w{s}*(PL!{FYC[2028]}{PB[s]['bv']}+PL!{FYC[2029]}{PB[s]['bv']})/2" for s in SC), None, None], EF,
     "同業倍數＝現價 ÷ 約 14 個月後（約 2027 年底）每股淨值；對應 12 個月後（約 2027-10）用約 2028 年底每股淨值＝（FY2028 年底＋FY2029 年底）÷ 2（會計年度 6 月，依月份調整）")
tset(q + 2, "comptp", "Comparable 目標價", [f"=B{q}*B{q + 1}", None, None], MF, "同業中位數 × 前瞻每股淨值")
fwd = {2028: "8/12", 2029: "4/12"}
feps = lambda c: "(" + "+".join(f"{w}*{c}{EPSR[y]}" for y, w in fwd.items()) + ")"
tset(q + 3, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=tp_wA*" + feps("B") + "+tp_wB*" + feps("C") + "+tp_wC*" + feps("D"), None, None], EF,
     "約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 8／12 × FY2028（2027-07～2028-06）＋4／12 × FY2029（會計年度 6 月，依月份調整）")
q += 5
V.cell(q, 1, "並列參考：同業前瞻 P/E（不進目標價）").font = BOLD; T["h02b"] = q
hdr(V, q + 1, ["同業與前瞻 P/E", "代號", "前瞻 P/E", "日期", "來源", "標記"]); T["ehd"] = q + 1
PEERS_E = [("美光", "MU", "=877.55/173", "2026-10-23", "現價 877.55 ÷ MarketScreener FY2027 共識 EPS 173（micron/ 模型 In 頁）"),
           ("SK 海力士", "000660.KS", "=1277000/475721", "2026-10-23", "現價 ÷ MarketScreener 2027 共識 EPS（skhynix/ 模型 In 頁）"),
           ("三星電子", "005930.KS", "=215500/72279", "2026-10-23", "現價 ÷ MarketScreener 2027 共識 EPS（samsung/ 模型 In 頁）"),
           ("Kioxia", "285A.T", "=12755/3486", "2026-10-23", "現價 ¥12,755 ÷ MarketScreener FY2027（2027-03）共識 EPS ¥3,486"),
           ("Western Digital", "WDC", 11.91, "2026-10-23", "stockanalysis 統計頁 Forward P/E（硬碟為主，參考）")]
for i, (nm, tk, pe, dt, src) in enumerate(PEERS_E, q + 2):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.number_format = "0.00"
    if not isinstance(pe, str): c.font = BLUE
    V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
T["e1"] = q + 2; T["el"] = q + 1 + len(PEERS_E)
q = T["el"] + 1
tset(q, "emed", "同業前瞻 P/E 中位數（並列參考）", [f"=MEDIAN(C{T['e1']}:C{T['el']})", None, None], "0.00")
tset(q + 1, "ecomp", "Comparable（P/E 並列參考）＝中位數 × 12 個月後前瞻 EPS", [f"=B{q}*B{T['feps']}", None, None], MF, "不進目標價：高峰期 EPS × 低倍數，與 P/B 法互相對照")
q += 3
V.cell(q, 1, "0.3 結論").font = BOLD; T["h03"] = q
hdr(V, q + 1, ["項目", "數值", "", "", "說明"]); T["hd3"] = q + 1
q += 2
tset(q, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], MF, "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁")
tset(q + 1, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], MF, "引用 0.2（前瞻 P/B）")
tset(q + 2, "cons", "市場共識目標價", ["=cons_tp", None, None], MF, "stockanalysis（S&P Global，25 位，2026-10-18 更新）：平均 2,144（中位數 2,150；區間 1,000–3,600）")
tset(q + 3, "upw", "上檔空間：加權目標價", [f"=B{q}/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-23 收盤）")
tset(q + 4, "upc", "上檔空間：Comparable 目標價", [f"=B{q + 1}/px-1", None, None], "0.0%")
x = f"B{q + 3}"
tset(q + 5, "tier", "上檔空間級距", [f'=IF({x}>0.5,"大於 50%",IF({x}>0.25,"25%–50%",IF({x}>=0,"0%–25%",IF({x}>=-0.25,"0%～−25%",IF({x}>=-0.5,"−25%～−50%","小於 −50%")))))', None, None], "@", "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
tset(q + 6, "rec", "投資建議", [f'=IF({x}>0.5,"強力買進",IF({x}>0.25,"買進",IF({x}>=-0.25,"中立",IF({x}>=-0.5,"放空","強力放空"))))', None, None], "@", "依上檔空間固定規則（Andy 2026-10-25 14:17）：>+50% 強力買進、+25～+50% 買進、−25～+25% 中立、−50～−25% 放空、<−50% 強力放空")
tset(q + 7, "dcfpx", "DCF ÷ 現價", [f"={c}{T['dcf12']}/px" for c in "BCD"], '0.00"x"', "各情境 DCF 價值（推到 12 個月後）÷ 現價（A／B／C）")
tset(q + 8, "gbv", "與市場共識目標價的差距：EPS 因素（記憶體以每股淨值計）", [f"=B{T['fbv']}/((cons_bv28+cons_bv29)/2)-1", None, None], "0.0%",
     "本模型約 2028 年底每股淨值（情境加權）÷ 共識（FY2028、FY2029 年底平均 601.9）− 1（Comparable 用 P/B，故以每股淨值代 EPS）")
tset(q + 9, "geps", "  對照：前瞻 EPS ÷ 共識前瞻 EPS − 1", [f"=B{T['feps']}/(8/12*cons_eps28+4/12*cons_eps29)-1", None, None], "0.0%", "共識前瞻 EPS＝8／12 × FY2028 267.6＋4／12 × FY2029 281.7（MarketScreener）")
tset(q + 10, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/((cons_bv28+cons_bv29)/2))/B{T['pmed']}-1", None, None], "0.0%")
V[f"E{q + 10}"] = f'="共識目標價隱含 P/B "&TEXT(cons_tp/((cons_bv28+cons_bv29)/2),"0.00")&"x（÷ 共識約 2028 年底每股淨值），同業中位數 "&TEXT(B{T["pmed"]},"0.00")&"x"'
tset(q + 11, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
T["meth"] = q + 13; T["obs"] = q + 14
V.cell(q + 13, 1, "目標價方法")
T["c_end"] = q + 13

# --- 第 5、6 節（第 80 列起） ---
b5 = 80
V.cell(b5, 1, "5. 目前倍數（僅供參考）").font = BOLD; hdr(V, b5 + 1, ["項目", "A／數值", "B", "C", "備註"]); CMH = b5 + 1
fr = b5 + 2; CUR = {}
for key, nm, fs, fm, note in (("px", "現價（2026-10-23 收盤）", ["=px"], EF, "stockanalysis、MarketScreener"),
                              ("pe26", "P/E：現價 ÷ FY2026 EPS 73.76", ["=px/a26_eps"], "0.0", "落後"),
                              ("pb26", "P/B：現價 ÷ FY2026 年底每股淨值", [f"=px/PL!D{PB['A']['bv']}"], "0.00", "FY2026 年底權益 15.736 ÷ 146.42 百萬股"),
                              ("pe27", "現價 ÷ FY2027 EPS（模型）", [f"=px/PL!E{SUMR[(s, 'eps')]}" for s in SC], "0.0", "A／B／C"),
                              ("pe28", "現價 ÷ FY2028 EPS（模型）", [f"=px/PL!F{SUMR[(s, 'eps')]}" for s in SC], "0.0", "A／B／C"),
                              ("pb28", "現價 ÷ FY2028 年底每股淨值（模型）", [f"=px/PL!F{SUMR[(s, 'bv')]}" for s in SC], "0.00", "A／B／C；第 0 節歷史分位檢查用")):
    V.cell(fr, 1, nm)
    for j, f in enumerate(fs): V.cell(fr, 2 + j, f).number_format = fm
    V.cell(fr, 5, note); CUR[key] = fr; fr += 1
def _y(s, y, k): return f"=PL!{FYC[y]}{PB[s][k]}"
def _cases(bp, bt, lp, lt):
    return [{"case": "見頂", "pos": "見頂年 → 25 分位", "q": 25, "basis_label": lp, "basis": bp},
            {"case": "中段", "pos": "中段 → 中位數", "q": 50, "basis_label": lp, "basis": bp},
            {"case": "谷底", "pos": "谷底 → 75 分位", "q": 75, "basis_label": lt, "basis": bt}]
vseries = [{"id": "P/B", "name": "P/B（主）", "sources": [{"label": "美光（Analogy；會計年度 8 月底）", "vals": MU_PB, "tag": "Analogy", "note": "來源：" + MU_PB_SRC + "（macrotrends）；SanDisk 自身上市不足十年"}],
            "cases": _cases(_y("A", 2027, "bv"), _y("C", 2028, "bv"), "FY2027 年底每股淨值（A）", "FY2028 年底每股淨值（C）")},
           {"id": "P/E", "name": "P/E（對照）", "sources": [{"label": "美光（Analogy；SEC 10-K 換算）", "vals": MU_PE, "tag": "Analogy", "note": "來源：" + MU_PE_SRC + "；虧損年不計"}],
            "cases": _cases(_y("A", 2027, "eps"), _y("C", 2028, "eps"), "FY2027 EPS（A）", "FY2028 EPS（C）")}]
nr = fr + 2
tr = VR.add_rule_block(V, nr, vseries, EF, "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底倍數依循環位置取分位。SanDisk 無十年自身歷史，以美光為 Analogy（同為記憶體循環股；差異：美光 DRAM 為主）。見頂與中段用 FY2027 A、谷底用 FY2028 C。")
PB0 = nr + 5; PB10 = f"$B${PB0}:$B${PB0 + 9}"
assert V.cell(PB0, 1).value == "2016" and V.cell(PB0 + 9, 1).value == "2025"
# 歷史分位（第 0 節）
hq = T["band"] + 1
tset(hq, "pr", "歷史倍數檢查：現價隱含 P/B 的十年分位（A／B／C）",
     [f"=IF(px/PL!F{SUMR[(s, 'bv')]}>=MAX({PB10}),1,IF(px/PL!F{SUMR[(s, 'bv')]}<=MIN({PB10}),0,PERCENTRANK({PB10},px/PL!F{SUMR[(s, 'bv')]})))" for s in SC], "0%",
     "現價 ÷ 各情境 FY2028 年底每股淨值，在 2016–2025 年底 P/B（第 6 節，美光 Analogy）分布的分位（高於最大值記 100%、低於最小值記 0%）；只作檢查，不進目標價（記憶體主估值為 P/B，故以 P/B 代 P/E）")
V.cell(T["meth"], 2, "DCF：股東現金流 FY2027–FY2029 用每股（淨利＋折舊－資本支出；資本支出含合資群分攤，約營收 6%），FY2030–FY2032 收斂期用 EPS ×（1 − g／ROE），r 10%、g 4%、ROE 為 FY2025–FY2026 平均；明確期＋永續，基準日 2026 年底。"
                       "Comparable：美光、SK 海力士、三星、Kioxia 前瞻 P/B 中位數 × 約 2028 年底每股淨值（情境加權）；P/E 並列參考不進目標價。"
                       "偏誤：FY2027–FY2029 為高峰期盈餘、收斂期淨利率以美光十年平均代替（SanDisk 無自身歷史）；大量買回壓低帳面淨值，P/B 法偏低估。")
# 舊法
o0 = tr
V.cell(o0, 1, "舊法：規則值情境加權（參考）").font = BOLD; T["old"] = o0
hdr(V, o0 + 1, ["項目", "數值", "", "", "說明"])
AD = {V.cell(i, 1).value: i for i in range(nr, tr) if isinstance(V.cell(i, 1).value, str) and V.cell(i, 1).value.startswith("採用 ")}
t1 = o0 + 2
V.cell(t1, 1, "規則值情境加權價（參考，非目標價）"); V.cell(t1, 2, f"=tp_wA*B{t1 + 2}+tp_wB*B{t1 + 3}+tp_wC*B{t1 + 4}").number_format = MF
V.cell(t1, 5, "權重 × 各情境隱含價；約為年底價格、未折現")
V.cell(t1 + 1, 1, "相對現價（規則值情境加權價 ÷ 現價 − 1）"); V.cell(t1 + 1, 2, f"=B{t1}/px-1").number_format = "0.0%"
for k_, (s_, lab, q_) in enumerate((("A", "情境 A（順利）：75 分位 × FY2027 年底每股淨值（A）", "75 分位"), ("B", "情境 B（平平）：中位數 × FY2027 年底每股淨值（B）", "中位數"), ("C", "情境 C（不如預期）：25 分位 × FY2027 年底每股淨值（C）", "25 分位"))):
    src = AD[f"採用 {q_}（P/B）"]
    V.cell(t1 + 2 + k_, 1, lab); V.cell(t1 + 2 + k_, 2, f"=B{src}*PL!E{SUMR[(s_, 'bv')]}").number_format = MF
V.cell(t1 + 5, 1, "舊法說明"); V.cell(t1 + 5, 2, "十年（2016–2025 年底）P/B 分位（美光 Analogy）× 各情境 FY2027 年底每股淨值：A＝75 分位、B＝中位數、C＝25 分位。第 6 節見頂／中段／谷底規則值供 AI 投資地圖抽數。")
T["old_end"] = t1 + 5
for k in ("dcftp", "comptp", "tpw"): V[f"B{T[k]}"].fill = KEY
assert T["obs"] < b5 - 1
V.column_dimensions["A"].width = 56
for c in "BCDFGH": V.column_dimensions[c].width = 15
V.column_dimensions["E"].width = 70

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A）：對 FY2027、FY2028 營收與 EPS 的影響。每個案例為一個完整推導區塊（本頁下方，與 PL 同一引擎）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收 $B", "FY2027 EPS", "相對基準", "FY2028 營收 $B", "FY2028 EPS", "相對基準", "FY2027 毛利率"])
CASES = [("市場層營收 −20%（價格；FQ2 FY27 起）", "價格係數 0.8", dict(pxm="0.8")), ("市場層營收 +20%（價格）", "價格係數 1.2", dict(pxm="1.2")),
         ("份額 ×0.9（營收與位元）", "份額係數 0.9", dict(shm="0.9")), ("份額 ×1.1", "份額係數 1.1", dict(shm="1.1")),
         ("每位元成本年降 10%（C 的值）", "cd_C", dict(cd="cd_C")), ("每位元成本年降 18%", "cd_A 高", dict(cd="cd_A_hi")),
         ("指引毛利率 83.0%（校準）", "g_gm 低", dict(gm="g_gm_lo")), ("指引毛利率 84.9%", "g_gm 高", dict(gm="g_gm_hi")),
         ("營業費用變動率 2%", "ox_var 高", dict(oxv="ox_var_hi")), ("稅率 20%", "tax 高", dict(tax="tax_hi")), ("稅率 12%", "tax 低", dict(tax="tax_lo")),
         ("買回 ÷ FCF 80%", "payout 高", dict(pay="payout_hi")), ("買回 ÷ FCF 30%", "payout 低", dict(pay="payout_lo"))]
def put(i, nm, st, R, pre=""):
    SE.cell(i, 1, nm); SE.cell(i, 2, st)
    SE[f"C{i}"] = f"={pre}E{R['yrev']}"; SE[f"D{i}"] = f"={pre}E{R['eps']}"; SE[f"F{i}"] = f"={pre}F{R['yrev']}"; SE[f"G{i}"] = f"={pre}F{R['eps']}"
    SE[f"E{i}"] = f"=D{i}/D$4-1"; SE[f"H{i}"] = f"=G{i}/G$4-1"; SE[f"I{i}"] = f"={pre}E{R['ygm']}"
    for c in "CF": SE[f"{c}{i}"].number_format = "#,##0.00"
    for c in "DG": SE[f"{c}{i}"].number_format = EF
    for c in "EHI": SE[f"{c}{i}"].number_format = "0.0%"
put(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base = 60
for n, (nm, st, ov) in enumerate(CASES):
    Rn = engine(SE, base + n * 50, f"案例區塊：{nm}（情境 A）", "A", **ov)
    put(5 + n, nm, st, Rn)
SN = 5 + len(CASES)
put(SN, "情境 B", "markets/nand B；成本年降 12%", PB["B"], "PL!"); put(SN + 1, "情境 C", "markets/nand C；成本年降 10%", PB["C"], "PL!")
SENS_LAST = SN + 1
SE.cell(SN + 3, 1, "讀法：FQ1 FY27 營收由市場層等於公司指引（恆等式），價格係數只作用在 FQ2 FY27 起；毛利率校準改變單位成本起點，之後每年依成本年降。")
assert SN + 3 < base
s0 = SE.max_row + 3
SE.cell(s0, 1, "目標價參數（只改 DCF 參數與權重；PL 收斂法的長期成長不隨 tp_g 變動）").font = BOLD
hdr(SE, s0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def dcf_formula(c, r_="tp_r", g="tp_g"):
    cf = lambda y: (f"Val!{c}{CFR[y]}" if y <= 2029 else f"Val!{c}{EPSR[y]}*(1-MIN({g}/tp_roe,0.8))")
    pv = "+".join(f"{cf(y)}/(1+{r_})^{t}" for t, y in enumerate(YRS6, 1))
    return f"MAX(0,{pv}+{cf(2032)}*(1+{g})/({r_}-{g})/(1+{r_})^6)"
tcases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r_="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r_="(tp_r-0.01)")),
          ("長期成長 +1pt", "g 5%（永續與再投資率）", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
          ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(tcases):
    i = s0 + 2 + k; w = kw.pop("w", "tp_w_dcf"); rr_ = kw.get("r_", "tp_r")
    SE.cell(i, 1, nm); SE.cell(i, 2, st)
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, **kw)}" for s, c in zip(SC, "BCD")) + f")*(1+{rr_})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"; SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"; SE[f"F{i}"] = f"=E{i}/E${s0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = MF
    SE[f"F{i}"].number_format = "0.0%"
SE.cell(s0 + 2 + len(tcases), 1, "核對：基準列 DCF 目標價 − Val 第 0 節（應為 0）"); SE[f"C{s0 + 2 + len(tcases)}"] = f"=C{s0 + 2}-Val!B{T['dcftp']}"
T["sens0"] = s0; T["sens_last"] = s0 + 1 + len(tcases)
SE.column_dimensions["A"].width = 40; SE.column_dimensions["B"].width = 30
for c in "CDEFGHI": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司實績、指引與共識（共識只作對照，未用於校準）。差距＝模型 ÷ 對照 − 1（百分點列為相減）"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rcr = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False):
    global rcr
    RC.cell(rcr, 1, label); RC.cell(rcr, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv = 3 + 2 * j; RC.cell(rcr, cv, getter(s)).number_format = fmt
        L_ = GL(cv); RC.cell(rcr, cv + 1, f"={L_}{rcr}-B{rcr}" if pt else f"={L_}{rcr}/B{rcr}-1").number_format = "0.0%"
    RC.cell(rcr, 9, src); RC.cell(rcr, 10, note); rcr += 1
RC.cell(rcr, 1, "A. 實績與公司指引").font = BOLD; rcr += 1
rec("FY2026 營收（實績重現）", "=a26_rev", lambda s: f"=PL!D{SUMR[(s, 'yrev')]}", Q4, "恆等式")
rec("FQ1 FY27 營收 vs 公司指引中值", "=g_rev", lambda s: f"=PL!{QC[(2027, 1)]}{PB[s]['rev']}", Q4, "恆等式（市場層份額由指引推得）")
rec("FQ1 FY27 毛利率 vs 指引中值（百分點）", "=g_gm", lambda s: f"=PL!{QC[(2027, 1)]}{PB[s]['gm']}", Q4, "恆等式（單位成本以此校準）", "0.0%", True)
rec("FQ1 FY27 EPS vs 指引中值（non-GAAP）", "=g_eps", lambda s: f"=PL!{QC[(2027, 1)]}{PB[s]['epsq']}", Q4, "本模型為 GAAP 營業費用＋15% 稅率，近似 non-GAAP（差約 SBC）")
RC.cell(rcr, 1, "B. 共識（MarketScreener 2026-10-23；只作參考）").font = BOLD; rcr += 1
rec("FY2027 營收 vs 共識", "=cons_rev27", lambda s: f"=PL!E{SUMR[(s, 'yrev')]}", MS)
rec("FY2028 營收 vs 共識", "=cons_rev28", lambda s: f"=PL!F{SUMR[(s, 'yrev')]}", MS)
rec("FY2027 EPS vs 共識", "=cons_eps27", lambda s: f"=PL!E{SUMR[(s, 'eps')]}", MS, "差距超過 30% 見 README", EF)
rec("FY2028 EPS vs 共識", "=cons_eps28", lambda s: f"=PL!F{SUMR[(s, 'eps')]}", MS, "", EF)
rec("FY2028 年底每股淨值 vs 共識", "=cons_bv28", lambda s: f"=PL!F{SUMR[(s, 'bv')]}", MS, "", EF)
RC.cell(rcr, 1, "C. 市場層一致性").font = BOLD; rcr += 1
RC.cell(rcr, 1, "SanDisk 營收 ÷ 市場層（公司口徑係數 sd_k）"); RC.cell(rcr, 3, "=sd_k").number_format = "0.0000"; RC.cell(rcr, 9, "In 頁"); rcr += 1
RC.cell(rcr, 1, "Kioxia＋SanDisk 合資群份額 CY2027（市場層 Ann 頁，A／B／C）")
for j, s in enumerate(SC): RC.cell(rcr, 3 + 2 * j, JVSH[s][0][2]).number_format = "0.0%"
RC.cell(rcr, 9, f"{rel(NAND_F)}｜Ann!F{JVSH['A'][1]}（kioxia/ 模型讀同一組列）"); rcr += 1
RC_LAST = rcr - 1
RC.column_dimensions["A"].width = 52
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 50; RC.column_dimensions["J"].width = 50

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源"; SO["A1"].font = BOLD
srcs = [("SanDisk FQ4 FY26 新聞稿（8-K，2026-08-05）", "https://www.sec.gov/Archives/edgar/data/0002023554/000162828026053346/sndkq4-26ex991xpressrelease.htm"),
        ("SanDisk FQ2 FY26 新聞稿（8-K，2026-01-29）", "https://www.sec.gov/Archives/edgar/data/2023554/000162828026004121/sndkq2fy26ex991-pressrelea.htm"),
        ("SanDisk：FQ1 FY27 財報日 2026-10-29（2026-09-29 公告）", "https://investor.sandisk.com/news-releases/news-release-details/sandisk-report-first-quarter-fiscal-year-2027-results-october-29"),
        ("stockanalysis：歷史股價、統計、共識目標價", "https://stockanalysis.com/stocks/sndk/forecast/"),
        ("MarketScreener：SanDisk 預估（EPS、每股淨值、股數）", "https://uk.marketscreener.com/quote/stock/SANDISK-CORPORATION-182801077/finances/"),
        ("MarketScreener：Kioxia 預估（同業 P/B、P/E）", "https://uk.marketscreener.com/quote/stock/KIOXIA-HOLDINGS-CORPORATI-180030735/finances/"),
        ("stockanalysis：WDC 統計（Forward P/E）", "https://stockanalysis.com/stocks/wdc/statistics/"),
        ("Motley Fool：SanDisk／Kioxia 日本投資、FY27 毛資本支出約營收 6%（2026-08-29）", "https://www.fool.com/investing/2026/08/29/sandisk-and-kioxia-plan-to-invest-more-than-usd31-billion-in-japanese-memory-plants-about-60-of-what-they-have-spent-there-in-25-years/"),
        (f"市場層：{rel(NAND_F)}（markets/nand v0.1）", "本 repo"),
        (f"美光十年倍數與淨利率（Analogy）：{rel(MU_F)}", "本 repo"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 80; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
VA = lambda k, c="B": f"Val!{c}{T[k]}"
OBS = [
 f'="1. 加權目標價 "&TEXT({VA("tpw")},"#,##0")&" 美元，相對現價 "&TEXT(px,"#,##0.00")&"（2026-10-23 收盤）"&TEXT({VA("upw")},"+0.0%;−0.0%")&"，主因是 Comparable 偏低：同業前瞻 P/B 中位數 "&TEXT({VA("pmed")},"0.00")&"x，而共識目標價隱含約 "&TEXT(cons_tp/((cons_bv28+cons_bv29)/2),"0.00")&"x（倍數因素 "&TEXT({VA("gmul")},"+0%;−0%")&"）；DCF 目標價 "&TEXT({VA("dcftp")},"#,##0")&" 高於 Comparable "&TEXT({VA("comptp")},"#,##0")&"。"',
 f'="2. 本模型 FY2027 EPS（A）"&TEXT(PL!E{SUMR[("A", "eps")]},"0")&" 美元，比共識 217.4 "&TEXT(PL!E{SUMR[("A", "eps")]}/cons_eps27-1,"+0%;−0%")&"；B、C 到 FY2028 才分歧（"&TEXT(PL!F{SUMR[("B", "eps")]},"0")&"／"&TEXT(PL!F{SUMR[("C", "eps")]},"0")&"），因為 2026Q3–2027Q2 的價格與位元三情境幾乎相同（markets/nand）。追蹤：2026-10-29 FQ1 FY27 財報與 FQ2 指引。"',
 f'="3. 最敏感的是價格而非成本：2027 起季價 ±20% 使 FY2028 EPS（A）變動 "&TEXT(Sens!H5,"+0%;−0%")&"／"&TEXT(Sens!H6,"+0%;−0%")&"，成本年降 ±4pt 只有約 ±1%；情境 C 的 FY2029 毛利率降到 "&TEXT(PL!G{SUMR[("C", "ygm")]},"0%")&"（CY2028 NAND 價格跌回約 2025 年水準）。追蹤：TrendForce 每季初的 NAND 合約價與 Kioxia 2026-10-30 財報（合資群位元與價格）。"']
PROP = ("SanDisk 是純 NAND 公司，盈餘由 NAND 合約價循環與 AI 推論帶動的企業級 SSD 位元決定，產能與 Kioxia 合資（Flash Ventures 49.9%）；"
        "問題是 2026 年的高價能維持多久（FY2028 起三情境分歧），以及現價隱含的倍數是否已反映循環見頂。")
lines = [(f"{NAME}（{TICKER} US）收支與估值模型 v0.1（AI 半導體 Project：記憶體——NAND）　2026-10-25", True),
         ("模型命題：" + PROP, False),
         ("重點觀察", True), (OBS[0], False), (OBS[1], False), (OBS[2], False),
         (f'="投資建議："&{VA("rec")}&"（依上檔空間固定規則；上檔空間："&TEXT({VA("upw")},"+0.0%;−0.0%")&"，級距："&{VA("tier")}&"）"', True), ("", False),
         ("驅動因子 → 推導量（因果方向：左邊決定右邊）", True),
         ("Units｜markets/nand 成品：SanDisk 季度營收（市場營收 × 份額；同字母情境）與位元（季度出貨 × 份額）", False),
         ("PL 區塊｜營收＝市場層 × 公司口徑係數；營業成本＝位元 × 每 EB 成本（FQ1 FY27 以指引毛利率校準，之後每年 × (1 − 成本年降)）；營業費用＝固定（指引校準）＋變動", False),
         ("PL 區塊｜淨利＝（營業利益＋淨現金收益）×（1 − 稅率）；自由現金流＝淨利＋折舊 − 資本支出（含合資群分攤）；買回＝回饋比例 × FCF → 股數 → EPS、每股淨值", False),
         ("Val｜第 0 節目標價＝DCF（FY2027–FY2032＋永續）與 Comparable（同業前瞻 P/B × 約 2028 年底每股淨值）各半；第 5–6 節目前倍數、十年倍數規則值（美光 Analogy）", False), ("", False),
         ("會計年度：約 6 月底結束（FY2027＝2026-07～2027-06 ≈ CY2026Q3–CY2027Q2）。前瞻 EPS 依月份調整：12 個月後（約 2027-10）的未來 12 月＝8／12 FY2028＋4／12 FY2029。", False),
         ("情境：A＝順利、B＝平平、C＝不如預期，讀 markets/nand 同字母（價格＝markets/memory NAND_S；顆數＝SCEN_UNITS；合資群份額年變動）；公司特有變數（每位元成本年降 14%／12%／10%）A 最有利、C 最不利。FY2027 的 B、C EPS 略高於 A：上游 markets/memory 的 B、C 價格路徑 2027 上半年漲幅較大、下半年才回落（行業情境，不在公司層調整）。FY2030 起為收斂法，非驅動因子推導。", False),
         ("十年歷史不足：2025-02 自 WDC 分拆上市，十年 P/B、P/E 與十年平均淨利率以美光為 Analogy。本版未做先驗，無 Andy 合理 P/E 對照。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 200
README_ROWS = {"thesis": 2, "obs": [4, 5, 6], "rec": 7, "drivers": [10, 11, 12, 13], "limits": 17}
assert RD.cell(17, 1).value.startswith("十年歷史不足") and RD.cell(9, 1).value.startswith("驅動因子")

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
RULER = [i for i in range(nr, tr) if isinstance(V.cell(i, 1).value, str) and V.cell(i, 1).value.startswith("規則值：")]
RULEH = RULER[0] - 1
RC_FMT = {str(i): {c: "p1" for c in "BCEG"} for i in range(4, RC_LAST + 1) if isinstance(RC.cell(i, 1).value, str) and "毛利率" in RC.cell(i, 1).value}
RC_FMT.update({str(i): {"C": "n2", "E": "n2", "G": "n2"} for i in range(4, RC_LAST + 1) if isinstance(RC.cell(i, 1).value, str) and "合資群份額" in RC.cell(i, 1).value})
for i in RC_FMT:
    if "合資群份額" in str(RC.cell(int(i), 1).value): RC_FMT[i] = {"C": "p1", "E": "p1", "G": "p1"}
EXTRA = dict(
    DRIVERS=[("g_rev", "n2", ["FY2027 營收（情境 A）：", f"@PL!E{SUMR[('A', 'yrev')]}|n1"], "公司指引決定 SanDisk 在市場層的份額起點；之後營收隨市場層價格與位元同向"),
             ("g_gm", "p1", ["FY2027 毛利率（情境 A）：", f"@PL!E{SUMR[('A', 'ygm')]}|p1"], "以指引毛利率校準每 EB 成本，之後依成本年降；價格上升時毛利率同向上升"),
             ("cd_A", "p0", ["FY2028 毛利率（情境 A）：", f"@PL!F{SUMR[('A', 'ygm')]}|p1"], "每位元成本下降使毛利率上升（同向）"),
             ("tax", "p0", ["FY2027 淨利（情境 A）：", f"@PL!E{SUMR[('A', 'yni')]}|n1"], "稅率上升使淨利下降（反向）"),
             ("cap_r", "p0", ["FY2027 自由現金流（情境 A）：", f"@PL!E{SUMR[('A', 'fcf')]}|n1"], "資本支出比例上升使自由現金流與 DCF 下降（反向）"),
             ("payout", "p0", ["FY2028 EPS（情境 A）：", f"@PL!F{SUMR[('A', 'eps')]}|n2"], "買回減少股數、提高 EPS；同時壓低每股淨值（P/B 法偏低）")],
    QUESTION="NAND 合約價的高檔能維持多久，以及現價隱含的倍數是否已反映循環見頂",
    PL_INTRO="FY2026 為實績；FY2027 起由 NAND 市場層驅動，三情境到 FY2028 才明顯分歧（會計年度約六月底結束）。",
    TITLER={s: (56 if s == "A" else PB[SC[SC.index(s) - 1]]["end"] + 2) for s in SC},
    RULEH=RULEH, RULER=RULER, RC_FMT=RC_FMT, TORNADO="FY2028 EPS 相對基準的變動（情境 A，一次動一個輸入）")
if os.environ.get("SNDK_ROWS"):
    json.dump(dict(**EXTRA, IN=IN_ROW, SUMR={f"{a}|{b}": v for (a, b), v in SUMR.items()}, PB=PB, EXT=EXT, T=T, EPSR=EPSR, CFR=CFR, PL_CHK=PL_CHK, CUR=CUR, CMH=CMH,
                   SENS_LAST=SENS_LAST, RC_LAST=RC_LAST, NR=nr, TR=tr, PB0=PB0, UR={f"{a}|{b}": v for (a, b), v in UR.items()}, RD=README_ROWS, EXTC=EXTC, P0=p0),
              open(os.environ["SNDK_ROWS"], "w"), ensure_ascii=False)
print("saved", sys.argv[1])
