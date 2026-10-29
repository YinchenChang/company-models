"""builder 與 v0.5 遷移的結構性測試（不需 LibreOffice）。"""
import sys
from collections import Counter
from pathlib import Path

import pytest
import yaml

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "builder"))
from v05_map import build_map, load  # noqa: E402
from engine import Engine, current_model_path  # noqa: E402

EXPECT = yaml.safe_load((Path(__file__).parent / "scenarios.yaml").read_text(encoding="utf-8"))["workbook_expectations"]


@pytest.fixture(scope="module")
def R():
    return build_map()


def test_every_v05_leaf_has_a_destination(R):
    assert len(R.leaves) == EXPECT["v05_leaves"]
    assert [p for p in R.leaves if p not in R.dest] == []


def test_tagged_objects_count():
    n = 0

    def w(o):
        nonlocal n
        if isinstance(o, dict):
            n += "tag" in o
            [w(v) for v in o.values()]
        elif isinstance(o, list):
            [w(v) for v in o]
    w(load())
    assert n == EXPECT["v05_tagged_objects"]


def _same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-12 * max(abs(a), abs(b), 1)       # LibreOffice 存檔保留 15 位有效數字
    return a == b


def test_src_and_inputs_values_equal_json(R):
    """SRC_OAI／Inputs 內每個有值的列，Excel 具名範圍的值必須等於 builder 由 v0.5 JSON 讀出的值。"""
    eng = Engine(current_model_path())
    for r in R.src_rows + R.inp_rows:
        if r["value"] != "":
            if isinstance(r["value"], str) and r["value"].startswith("="):      # v0.6.4 W4：公式值的列另測（test_w4_*）
                continue
            assert _same(eng.get_name(r["id"]), r["value"]), r["id"]


def test_classification_rules(R):
    for r in R.src_rows:
        assert not str(r["tag"]).startswith(("Assumed", "Analogy")), r["id"]
    kinds = Counter(k for k, _, _ in R.dest.values())
    assert kinds["SRC"] > 0 and kinds["INP"] > 0
    for p, (k, ref, why) in R.dest.items():
        if k == "TK":
            assert "TK_Link" in ref


def test_pending_list_resolved(R):
    """r3 E8 已逐項處理 v0.5 分類不明的 A 類項目；日後新增不明項目才會出現在 pending。"""
    assert isinstance(R.pending, list) and R.pending == []


def _e6_scan(ws):
    import re
    seen = 0
    for row in ws.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                seen += 1
                f = re.sub(r"'?[A-Za-z_]+'?!\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?", "", c.value)    # 跨頁參照
                f = re.sub(r"\$?\b[A-Z]{1,3}\$?\d+\b", "", f)                          # 儲存格參照
                f = re.sub(r"\b[A-Za-z_][A-Za-z_0-9]*\b", "", f)                          # 具名範圍與函數
                f = f.replace("(1-", "(").replace("+1)", ")").replace("(1+", "(")   # 恆等式：1−比例；年數含頭尾的 +1；1＋成長率（P2）
                assert not re.search(r"\d", f), f"{ws.title}!{c.coordinate} 公式含常數：{c.value}"
    return seen


def test_e6_no_constants_in_v9_formulas():
    """E6：Derived_V9 的公式不得內含常數（算式中的『1-比例』恆等式除外）；換算係數與比例必須在 Inputs。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    assert _e6_scan(wb["Derived_V9"]) >= 9


def test_e6_no_constants_in_p2_formulas():
    """E6（P2）：Demand、Revenue 兩頁的公式不得內含常數；恆等式只容許 1−比例、1＋成長率；定義常數（天數、月數、單位換算）在 Inputs。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    assert _e6_scan(wb["Demand"]) >= 200
    assert _e6_scan(wb["Revenue"]) >= 200


def test_v12_pro_constraints_in_src(R):
    """V12：SRC_OAI 只登 pro 的兩個約束（倍數、占比上限），不再登 pro 的用戶人數點值。"""
    pro = [r for r in R.src_rows if r["metric"].startswith("Pro ")]
    assert len(pro) == 2
    assert {r["unit"] for r in pro} == {"倍", "比例"}


def test_e6_no_constants_in_p3_formulas():
    """E6（P3）：Compute 頁的公式不得內含常數；恆等式只容許 1−比例、年數 +1、1＋成長率；定義常數（T→M、等差級數除數）在 Inputs。
    世代名稱與成本情境為文字格（SUMIFS 的鍵），不在公式內。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    assert _e6_scan(wb["Compute"]) >= 600


def test_e6_no_constants_in_p4_formulas():
    """E6（P4）：Cost 頁的公式不得內含常數；恆等式只容許 1−比例、年數 +1、1＋成長率；定義常數（$M→$B、$→$B、等差級數除數）在 Inputs。
    世代名稱與成本情境為文字格（SUMIFS 的鍵），不在公式內。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    assert _e6_scan(wb["Cost"]) >= 500
    assert _e6_scan(wb["Compute"]) >= 620


def test_e6_no_constants_in_p5_formulas():
    """E6（P5）：Funding、Reverse 兩頁的公式不得內含常數；恆等式只容許 1−比例、年數 +1、1＋成長率；定義常數（$M→$B、月數）在 Inputs。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    assert _e6_scan(wb["Funding"]) >= 150
    assert _e6_scan(wb["Reverse"]) >= 180


def test_reverse_not_fed_back():
    """R1／共同規則第 4 節：反向模式只作對照區塊——除 Checks 外，沒有任何頁的公式引用 Reverse 頁或 RVS_ 具名範圍。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    for ws in wb.worksheets:
        if ws.title in ("Reverse", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "Reverse!" not in c.value and "RVS_" not in c.value, f"{ws.title}!{c.coordinate} 引用反向模式：{c.value}"


def test_e6_no_constants_in_pxq_formulas():
    """E6（v0.6.1）：PxQ 頁的公式不得內含常數；恆等式只容許 1−比例、年數 +1。文字格式字串（TEXT 的 "0.0"）與引號內文字不計。"""
    import re
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    ws = wb["PxQ"]
    seen = 0
    for row in ws.iter_rows(min_row=6):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                seen += 1
                f = re.sub(r'"[^"]*"', "", c.value)
                f = re.sub(r"'?[A-Za-z_]+'?!\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?", "", f)
                f = re.sub(r"\$?\b[A-Z]{1,3}\$?\d+\b", "", f)
                f = re.sub(r"\b[A-Za-z_][A-Za-z_0-9]*\b", "", f)
                f = f.replace("(1-", "(").replace("+1)", ")").replace("-1,", ",").replace("-1)", ")")
                f = re.sub(r"\+1$", "", f)                               # 年數含頭尾的 +1（年度序位）
                assert not re.search(r"\d", f), f"PxQ!{c.coordinate} 公式含常數：{c.value}"
    assert seen >= 700


def test_pxq_reference_whitelist():
    """v0.6.2：PxQ 頁公式只許引用 TK_（且須在 P×Q 取數名稱清單內，或表頭 TK_HdrGen／TK_HdrCost）、INP_、SRC_OAI_、CMP_、REV_ 與本頁名稱（PXQ_）；不得引用 B4_／AL_／L1_／CTL_，也不得直接引用其他頁儲存格。"""
    import re
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    ws = wb["PxQ"]
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "builder"))
    from tk_link import PXQ_NAMES
    tk_ok = {"TK_" + n for n in PXQ_NAMES} | {"TK_HdrGen", "TK_HdrCost"}
    allowed = ("TK_", "INP_", "SRC_OAI_", "CMP_", "REV_", "PXQ_")
    funcs = {"SUM", "SUMPRODUCT", "SUMIFS", "SUMSQ", "INDEX", "LN", "AVERAGE", "IF", "ABS", "TEXT", "COUNTIF"}
    for row in ws.iter_rows(min_row=6):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                f = re.sub(r'"[^"]*"', "", c.value)
                assert "!" not in f, f"PxQ!{c.coordinate} 直接引用其他頁：{c.value}"
                for tok in re.findall(r"\b[A-Za-z_][A-Za-z_0-9]*\b", f):
                    if tok in funcs or re.fullmatch(r"[A-Z]{1,3}\d+", tok) or re.fullmatch(r"[A-Z]{1,3}", tok):
                        continue
                    assert tok.startswith(allowed), f"PxQ!{c.coordinate} 引用不在白名單：{tok}（{c.value}）"
                    if tok.startswith("TK_"):
                        assert tok in tk_ok, f"PxQ!{c.coordinate} 引用 TK_ 名稱不在 P×Q 取數清單：{tok}"
                    assert not tok.startswith(("TK_B4_", "TK_AL_", "TK_L1_", "TK_CTL_")), tok


def test_pxq_not_fed_back():
    """PxQ 是下游對照頁：除 Checks 外，沒有任何頁的公式引用 PxQ 或 PXQ_ 名稱。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path())
    for ws in wb.worksheets:
        if ws.title in ("PxQ", "Checks"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "PxQ!" not in c.value and "PXQ_" not in c.value, f"{ws.title}!{c.coordinate}"


def test_tk_link_single_snapshot_v21_7():
    """v0.6.4：TK_Link 為單一 Tokenomics v21.7 快照（v0.6.2 起單一快照；v0.6.3 換版 v19.0→v21.4；v0.6.4 換版 v21.4→v21.7，88 名數值逐名不變）（88 名有值、待提供 2 名）；TK_PxQ 頁與 TKQ_ 名稱已退役；前 63 名列位與 v0.6.1 相同（名稱、列序不變）。"""
    import openpyxl
    root = Path(__file__).resolve().parents[2]
    old = openpyxl.load_workbook(root / "model" / "archive" / "20261022_OpenAI_v0.6.1.xlsx", data_only=True)["TK_Link"]
    new = openpyxl.load_workbook(current_model_path(), data_only=True)
    assert "TK_PxQ" not in new.sheetnames
    assert not any(n.startswith("TKQ_") for n in new.defined_names)
    cur = new["TK_Link"]
    assert cur["B4"].value == "v21.7"
    n = 0
    for r in range(10, 200):
        if old.cell(r, 1).value is None:
            continue
        n += 1
        assert old.cell(r, 1).value == cur.cell(r, 1).value, r          # 名稱、列序不變
        assert old.cell(r, 2).value == cur.cell(r, 2).value, r
    assert n == 63
    ok = sum(1 for r in range(10, 200) if cur.cell(r, 6).value == "OK")
    assert ok == 88
    for r in range(10, 200):
        v = cur.cell(r, 1).value
        assert not (v and str(v).startswith(("B4_", "AL_", "CTL_")))
        assert not (v and str(v).startswith("L1_") and not str(v).startswith("L1_Ans3"))   # L1_ 只許 L1_Ans3（Q3 對照；下游契約第 15 條）


def _pxq_values():
    """以 LibreOffice 重算後的值讀 PxQ（建置時已重算並存檔；此處讀模型檔內的快取值）。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path(), data_only=True)
    return wb["PxQ"]


def test_pxq_three_scope_structure():
    """v0.6.3：PxQ 頂部主輸出＝三口徑 × 低／基準／高；2025 對照；倍數與持平門檻表齊備；欄 D–N 共 11 欄案例。"""
    ws = _pxq_values()
    hdr = [ws.cell(5, c).value for c in range(4, 15)]
    assert hdr[:2] == ["2025 指數內插", "2025 線性內插"]
    assert [h for h in hdr[2:]] == [f"口徑{z} {lv}" for z in "一二三" for lv in ("低", "基準", "高")]
    base = {1: (9, 2.26, 26.3), 2: (10, 3.0, 35.0), 3: (11, 4.6, 50.0)}      # v0.6.4 W1：口徑一基準 GW 2.34→2.26
    for s, (r, gw, rev) in base.items():
        assert abs(ws.cell(r, 8).value - gw) < 1e-9 and abs(ws.cell(r, 10).value - rev) < 1e-9
    r25 = ws["F12"].value
    assert abs(r25 - 1.239) < 0.001
    for s, mult in ((1, 0.959), (2, 0.96), (3, 0.90)):
        assert abs(ws.cell(14 + s, 5).value - mult) < 0.01 * mult + 0.005, (s, ws.cell(14 + s, 5).value)
    # 持平門檻（A 段表）
    expect = {21: (2.4014, 2.1669, 1.9218), 22: (3.1958, 2.8837, 2.5575), 23: (4.5654, 4.1196, 3.6536)}
    for r, vals in expect.items():
        for c, v in zip((4, 5, 6), vals):
            assert abs(ws.cell(r, c).value / v - 1) < 0.01, (r, c)


def test_w3_gw_route_token_ledger():
    """W3：GW 路線 token 帳 (A) 交接檔組合×0.85 ＝2.10；(B) ×0.806 ＝1.99；容差 1%。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path(), data_only=True)
    ws = wb["PxQ"]
    names = {n: wb.defined_names[n].attr_text for n in ("PXQ_TokLedger", "PXQ_TokLedgerRel")}
    sheet, rng = names["PXQ_TokLedger"].split("!")
    cells = [c.value for c in ws[rng.replace("$", "")][0]]
    assert abs(cells[0] / 2.10 - 1) <= 0.01 and abs(cells[1] / 1.99 - 1) <= 0.01, cells


def test_fairwater_six_scenarios_match_j7():
    """J7：六組 Q2 GW（占比 × 日期）的倍數與證據檔一致（容差 1%），範圍 0.76–0.98。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path(), data_only=True)
    sheet, rng = wb.defined_names["PXQ_FW_Mult"].attr_text.split("!")
    got = [r[0].value for r in wb["PxQ"][rng.replace("$", "")]]
    exp = [0.98249, 0.90648, 0.92599, 0.83742, 0.85707, 0.76092]
    assert len(got) == 6
    for g, e in zip(got, exp):
        assert abs(g / e - 1) <= 0.01, (g, e)
    assert abs(min(got) - 0.76) < 0.01 and abs(max(got) - 0.98) < 0.01


def test_pxq_recon_6b_clean():
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path(), data_only=True)
    sheet, rng = wb.defined_names["PXQ_ReconBad"].attr_text.split("!")
    assert wb["PxQ"][rng.replace("$", "")].value == 0


def test_w1_scope1_gw_values():
    """v0.6.4 W1：口徑一 Q2 季平均 GW 基準／低／高＝2.26／2.17／2.35（INP_289–291）；J7 列 INP_297 不改（2.34）。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path(), data_only=True)
    ws = wb["Inputs"]
    got = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in (5, 6, 7)] for r in range(5, ws.max_row + 1)}
    assert got["INP_289"] == [2.26, 2.17, 2.35]
    assert got["INP_290"] == [2.17, 2.17, 2.17] and got["INP_291"] == [2.35, 2.35, 2.35]
    assert got["INP_297"][0] == 2.34
    px = wb["PxQ"]
    for c, exp in zip((4, 5, 6), (0.999, 0.959, 0.921)):          # B 段 README：倍數 0.999／0.959／0.921
        assert abs(px.cell(15, c).value - exp) < 0.01 * exp, (c, px.cell(15, c).value)


def test_w2_d2_derivation_matches_b_segment():
    """v0.6.4 W2：第 8.4 節（A 證據重現）對 B 段 token帳：第 5 列 3.531／5.071／6.610、第 6 列 7.606、token 帳 k＝1 為 1.857／1.769／1.689，k＝2 為 1.666／1.530／1.414（容差 1%）；對帳項數 0；基準 token 帳 (A)(B) 不變。"""
    import openpyxl
    wb = openpyxl.load_workbook(current_model_path(), data_only=True)
    ws = wb["PxQ"]
    rows = {str(ws.cell(r, 2).value).split(" ")[0]: r for r in range(290, ws.max_row + 1) if ws.cell(r, 1).value and str(ws.cell(r, 2).value)[:2] in ("1.", "2a", "2b", "3.", "4.", "5.", "5′", "6.", "7a", "7b", "7c", "7d")}
    def vals(key):
        return [ws.cell(rows[key], c).value for c in (4, 5, 6)]
    for key, exp in (("5.", (3.531, 5.071, 6.610)), ("5′.", (7.063, 10.141, 13.220)), ("7a.", (1.857, 1.769, 1.689)), ("7b.", (1.666, 1.530, 1.414)), ("3.", (0.3503, 0.5030, 0.6558))):
        for g, e in zip(vals(key), exp):
            assert abs(g / e - 1) <= 0.01, (key, g, e)
    assert abs(vals("6.")[1] / 7.606 - 1) <= 0.01
    sheet, rng = wb.defined_names["PXQ_D2ReconBad"].attr_text.split("!")
    assert ws[rng.replace("$", "")].value == 0
    led = [c.value for c in ws[wb.defined_names["PXQ_TokLedger"].attr_text.split("!")[1].replace("$", "")][0]]
    assert abs(led[0] / 2.0966 - 1) < 0.001 and abs(led[1] / 1.9876 - 1) < 0.001
    assert abs(led[2] / (1.9365 * 2.26 / 2.34) - 1) < 0.001       # (C)＝v0.6.3 值 × 2.26 ÷ 2.34


def test_w4_derived_revenue_inputs():
    """v0.6.4 W4：SRC_OAI_141–143 退役並移入 Inputs（公式）；基準值與 v0.6.3 相同（26.8／26.3／35）；SRC_OAI_144、145 不動。"""
    import json
    import openpyxl
    root = Path(__file__).resolve().parents[2]
    reg = json.loads((root / "builder" / "id_registry.json").read_text(encoding="utf-8"))
    assert {"SRC_OAI_141", "SRC_OAI_142", "SRC_OAI_143"} <= set(reg["retired"]["SRC"])
    wf = openpyxl.load_workbook(current_model_path())
    wv = openpyxl.load_workbook(current_model_path(), data_only=True)
    assert not any(n in wf.defined_names for n in ("SRC_OAI_141", "SRC_OAI_142", "SRC_OAI_143"))
    assert "SRC_OAI_144" in wf.defined_names and "SRC_OAI_145" in wf.defined_names
    want = {"營收：2026 Q2 年化（含廣告；對照）": 26.8, "營收：2026 Q2 年化（不含廣告）": 26.3, "營收：FY2026 預測（不含廣告）": 35.0}
    seen = 0
    for r in range(5, wf["Inputs"].max_row + 1):
        nme = wf["Inputs"].cell(r, 2).value
        if nme in want:
            seen += 1
            assert str(wf["Inputs"].cell(r, 5).value).startswith("="), nme
            assert abs(wv["Inputs"].cell(r, 5).value - want[nme]) < 1e-9, nme
    assert seen == 3
    for r in range(5, wf["SRC_OAI"].max_row + 1):
        assert wf["SRC_OAI"].cell(r, 1).value not in ("SRC_OAI_141", "SRC_OAI_142", "SRC_OAI_143")
