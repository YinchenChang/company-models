#!/usr/bin/env python3
"""v0.6.1 對照 Excel：①新增列 ②比值與中間量全表 ③已套用的預設。數值取自 model/CURRENT 的快取值（LibreOffice 重算）。
用法：python3 tools/make_pxq_report.py --out docs/reports/20261022_v0.6.1_PxQ_對照.xlsx"""
import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
from engine import current_model_path  # noqa: E402

DEFAULTS = [
    ("L 讀法（判斷題 1）", "壽命期價格係數：隱含 m＝實際 m×L ÷ L，L 取 TK_IF_LTotalFleetVDisc（2025 機隊）與 TK_IF_LTotalDisc（新世代）；只作顯示", "表列牌價另列為一因子（P_eff ÷ (1−d)）", "不影響比值；影響隱含 m 的讀數"),
    ("服務占比（判斷題 2）", "基準 TK_IF_ServeShare 0.41；對照 0.50（Epoch『研發與推論大致平分』，SRC_OAI_140）", "其他占比", "比值不變；Q 與 m×L 的水準隨之改變（0.50：2025 線性欄實際 m×L 1.51→1.24）"),
    ("2026 營收（判斷題 3）", "SRC_OAI_043 管理層目標 36（FT 2026-09-18 七月投資人簡報；改標為利害關係方的管理層目標）；對照 REV_GrossCapped 34.2", "模型正向 34.2", "基準 2026 比值 1.225→1.164（GW 3.0）"),
    ("世代組合（判斷題 4）", "基準＝v0.6 自有 V1（CMP_Share_*；自研＝VR200×INP_204）；對照＝交接檔 2026 Hopper .25／GB200 .45／GB300 .25／VR200 .05（INP_281–286）", "交接檔組合", "2026 h 9.80→9.73 $B/GW；比值 1.225→1.233"),
    ("2025 GW 主值（判斷題 5）", "指數內插 1.128 GW 為主值（INP_276＝0）；線性 1.25 並列", "線性內插", "2025 比值 1.239→1.118"),
    ("生產折減（判斷題 6）", "v19.0 讀法：VR200 用 TK_IF_FleetProdRatio_VR200（0.806），其他世代用 TK_IF_ProdDerate（0.85）；自研／其他、Rubin Ultra 用 ProdDerate", "交接檔一律 0.85（INP_287＝1）", "2026 含 VR200 占比 5%：Q 約 −0.6%（2026 基準欄 2.446→2.461 ×10¹⁰ M tok）；比值不變"),
    ("自研晶片", "持有成本與付費年產能＝VR200 × INP_204（0.6）", "—", "沿用 v0.6 P3／P4 做法"),
    ("2026 平均 GW", "3.0（2.5–3.5）三情境；電力口徑未明", "—", "比值反比於 GW"),
    ("持平所需 GW", "2026 營收 ÷（2026 h × 2025 比值）；指數／線性各算", "—", "主值依 INP_276"),
    ("預覽對帳容許誤差", "0.5%（INP_288）", "—", "—"),
    ("TK_Link 讀取日", "保持 2026-10-08（v5.26 快照未重讀；--tk-date）", "改 2026-10-22", "TK_Link 63 名逐格相同"),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args()
    m = current_model_path()
    wb = openpyxl.load_workbook(m, data_only=True)
    out = openpyxl.Workbook()
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="FFEDEDED")
    # ① 新增列
    ws = out.active
    ws.title = "①新增列"
    ws.append(["頁", "ID／位置", "名稱", "值", "低", "高", "單位", "標記", "來源／備註"])
    s = wb["SRC_OAI"]
    for r in range(5, s.max_row + 1):
        if s.cell(r, 1).value in ("SRC_OAI_138", "SRC_OAI_139", "SRC_OAI_140"):
            ws.append(["SRC_OAI", s.cell(r, 1).value, s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 5).value, s.cell(r, 6).value, s.cell(r, 17).value, s.cell(r, 10).value])
    s = wb["Inputs"]
    for r in range(5, s.max_row + 1):
        i = s.cell(r, 1).value
        if i and str(i).startswith("INP_") and int(i[4:]) >= 276:
            ws.append(["Inputs", i, s.cell(r, 2).value + (f"（{s.cell(r, 3).value}）" if s.cell(r, 3).value not in (None, "—") else ""), s.cell(r, 5).value, s.cell(r, 6).value, s.cell(r, 7).value, s.cell(r, 4).value, s.cell(r, 8).value, s.cell(r, 12).value])
    s = wb["Checks"]
    for r in range(5, s.max_row + 1):
        i = s.cell(r, 1).value
        if i and str(i).startswith("C") and str(i)[1:].isdigit() and int(i[1:]) >= 131:
            ws.append(["Checks", i, s.cell(r, 2).value, s.cell(r, 3).value, s.cell(r, 4).value, s.cell(r, 5).value, "", "", s.cell(r, 6).value])
    for tab in ("PxQ",):
        ws.append([tab, "（新頁）", f"{wb[tab].max_row} 列 × {wb[tab].max_column} 欄", "", "", "", "", "", "見報告"])
    # ② 全表
    w2 = out.create_sheet("②比值與中間量")
    p = wb["PxQ"]
    w2.append(["編號", "項目", "單位", "2025 指數", "2025 線性", "2026 低 GW", "2026 基準 GW", "2026 高 GW", "說明"])
    for r in range(8, p.max_row + 1):
        row = [p.cell(r, c).value for c in (1, 2, 3, 4, 5, 6, 7, 8, 10)]
        if any(v is not None for v in row):
            w2.append(row)
    # ③ 預設
    w3 = out.create_sheet("③已套用的預設")
    w3.append(["問題", "採用的預設", "替代選項", "對結果的影響"])
    for d in DEFAULTS:
        w3.append(list(d))
    for w in (ws, w2, w3):
        for c in w[1]:
            c.font, c.fill = bold, fill
        for col in "ABCDEFGHI":
            w.column_dimensions[col].width = 28
        w.column_dimensions["B"].width = 60
        w.freeze_panes = "A2"
        for row in w.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
                c.alignment = Alignment(wrap_text=False, vertical="top")
    a.out.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
