#!/usr/bin/env python3
"""v0.6.2 TK 同步歸因：一次只換一組 Tokenomics 名稱，以 LibreOffice 重算整本活頁簿，取主要輸出（不在 Python 另寫計算）。

方法：以 v0.6.2 活頁簿為底，把 TK_Link 指定名稱組的值改回 v5.26（取自 v0.6.1 活頁簿 TK_Link）或保留 v19.0，存成變體檔，
LibreOffice 重算，從具名範圍讀輸出。情境：
  OLD    四組全為 v5.26（＝v0.6.1 的 TK_Link 值）
  NEW    四組全為 v19.0（＝v0.6.2）
  ONLY_g 只有 g 組為 v19.0（其餘 v5.26）
  LOO_g  只有 g 組為 v5.26（其餘 v19.0）
組別：1 IF_TokGW_*；2 IF_HoldEcon；3 IF_CapexTotal；4 其餘（兩版都有、值不同的名稱，且不在 1–3）。
v0.6.3（TK v19.0 → v21.4）加細：4 IF_TokGWPaidMix＋IF_FleetProdRatio_VR200（P×Q 的 Q 面）；5 IF_PxEffMix（P×Q 的牌價）；6 其餘（旗標 --split-q）。
用法：python3 tools/tk_attribution.py --new model/<新>.xlsx --old model/archive/<舊>.xlsx --out <json> [--split-q]
"""
import argparse
import json
import re
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "builder"))
from common import lo_recalc  # noqa: E402

OUTPUTS = {  # 名稱 → 說明；皆為具名範圍（年度向量取 FY2025–2030 六格；純量取一格）
    "CMP_Eta": "η（各年）", "CMP_Eta2025": "η 2025", "CMP_InfGW_Eff": "有效推論 GW（未截頂）", "CMP_InfGW": "推論 GW（截頂後）",
    "CMP_RDGW": "研發 GW", "CMP_DemandGW": "總需求 GW", "CMP_CapFactor": "截頂係數", "CMP_SupplyGW": "供給 GW", "CMP_Supply_Owned": "自建 GW",
    "COST_SupplierHold": "供應商持有成本 $B", "COST_CloudGM": "雲端毛利 $B", "COST_CloudGMPct": "雲端毛利 %",
    "COST_PropGap_VR": "每 VR 等值 GW 差額（現金）", "COST_PropGapEcon_VR": "每 VR 等值 GW 差額（經濟）",
    "COST_Compute": "算力成本（合約實付＋自有）", "COST_ComputeEcon": "算力成本（經濟）", "FND_ExtNeed": "外部資金需求（年）", "FND_ExtNeedCum": "累計外部資金需求",
    "CMP_TokGW": "token 換算 GW（÷η 前）", "CMP_VReqFactor": "機隊 VR 等值係數", "CMP_Supply_VReq": "供給 VR 等值 GW",
    "COST_PropRev_VR": "每 VR 等值 GW 營收 $B", "COST_PropFull_VR": "每 VR 等值 GW 全成本（現金）$B", "COST_Coverage": "營收覆蓋全成本比",
    "COST_RDCompute": "研發算力成本 $B", "COST_Q3": "Q3 研發占整體成本", "COST_Q3GapTK": "研發算力 − TK L1_Ans3 差距 $B",
    "PXQ_Ratio": "PxQ 比值", "PXQ_Q": "PxQ Q（M tok/年）", "PXQ_Peff": "PxQ 有效牌價", "PXQ_h": "PxQ 每 GW 持有成本 h", "PXQ_C": "PxQ 算力成本 C",
    "PXQ_mL_Actual": "PxQ 實際 m×L", "PXQ_mL_Required": "PxQ 所需 m×L", "CHK_Errors": "Checks 錯誤數",
}


def values(wb, name):
    t = wb.defined_names[name].attr_text
    sh, ref = t.rsplit("!", 1)
    ws = wb[sh.strip("'")]
    cells = ws[ref.replace("$", "")]
    if not isinstance(cells, tuple):
        return [cells.value]
    return [c.value for row in cells for c in (row if isinstance(row, tuple) else (row,))]


def tk_rows(ws):
    out = {}
    for r in range(10, 300):
        n = ws.cell(r, 1).value
        if n:
            out[n] = r
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--new", type=Path, required=True)
    ap.add_argument("--old", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--split-q", action="store_true", help="v0.6.3：組 4 拆為 Q 面（IF_TokGWPaidMix、IF_FleetProdRatio_VR200）、5 牌價（IF_PxEffMix）、6 其餘")
    a = ap.parse_args()
    new_ws = openpyxl.load_workbook(a.new, data_only=True)["TK_Link"]
    old_ws = openpyxl.load_workbook(a.old, data_only=True)["TK_Link"]
    nr, orr = tk_rows(new_ws), tk_rows(old_ws)
    diff = []
    for n, r in nr.items():
        if n in orr:
            nv = [new_ws.cell(r, c).value for c in range(10, 25)]
            ov = [old_ws.cell(orr[n], c).value for c in range(10, 25)]
            if nv != ov:
                diff.append(n)
    groups = {1: [n for n in diff if re.fullmatch(r"IF_TokGW_(Luna|Sol|Astra)", n)], 2: [n for n in diff if n == "IF_HoldEcon"],
              3: [n for n in diff if n == "IF_CapexTotal"]}
    groups[4] = [n for n in diff if n not in sum(groups.values(), [])]
    if a.split_q:
        q = [n for n in groups[4] if n in ("IF_TokGWPaidMix", "IF_FleetProdRatio_VR200")]
        px = [n for n in groups[4] if n == "IF_PxEffMix"]
        groups[5], groups[6] = px, [n for n in groups[4] if n not in q + px]
        groups[4] = q
    print({g: len(v) for g, v in groups.items()}, "異動名稱", len(diff))

    def variant(old_groups, tag):
        wb = openpyxl.load_workbook(a.new)            # 公式版
        ws = wb["TK_Link"]
        for g in old_groups:
            for n in groups[g]:
                for c in range(10, 25):
                    ws.cell(nr[n], c).value = old_ws.cell(orr[n], c).value
        d = Path(tempfile.mkdtemp())
        f = d / f"{tag}.xlsx"
        wb.save(f)
        rc = lo_recalc(f, d / "out")
        res = openpyxl.load_workbook(rc, data_only=True)
        return {k: values(res, k) for k in OUTPUTS}

    gs = sorted(groups)
    scen = {"OLD": variant(gs, "OLD"), "NEW": variant([], "NEW")}
    for g in gs:
        scen[f"ONLY_{g}"] = variant([x for x in gs if x != g], f"ONLY_{g}")
        scen[f"LOO_{g}"] = variant([g], f"LOO_{g}")
        print("done", g, flush=True)
    a.out.write_text(json.dumps(dict(outputs=OUTPUTS, groups=groups, scen=scen), ensure_ascii=False, indent=1, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
