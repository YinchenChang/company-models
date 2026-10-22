import openpyxl,sys,collections
a=openpyxl.load_workbook(sys.argv[1]); b=openpyxl.load_workbook(sys.argv[2])
av=openpyxl.load_workbook(sys.argv[1],data_only=True); bv=openpyxl.load_workbook(sys.argv[2],data_only=True)
print('sheets only in new:',[s for s in b.sheetnames if s not in a.sheetnames])
print('names only new:',len(set(b.defined_names)-set(a.defined_names)),'removed:',set(a.defined_names)-set(b.defined_names))
chg=[k for k in a.defined_names if k in b.defined_names and a.defined_names[k].attr_text!=b.defined_names[k].attr_text]; print('name ref changed:',chg)
for s in a.sheetnames:
    if s not in b.sheetnames: print('MISSING',s);continue
    wa,wb_=a[s],b[s]; va,vb=av[s],bv[s]
    d=collections.Counter(); ex=[]
    mr=max(wa.max_row,wb_.max_row); mc=max(wa.max_column,wb_.max_column)
    for r in range(1,mr+1):
        for c in range(1,mc+1):
            x,y=wa.cell(r,c).value,wb_.cell(r,c).value
            if x!=y:
                kind='formula' if (isinstance(x,str) and x.startswith('=')) or (isinstance(y,str) and str(y).startswith('=')) else 'const'
                d[kind]+=1; ex.append((wa.cell(r,c).coordinate,kind,str(x)[:50],str(y)[:50]))
            else:
                p,q=va.cell(r,c).value,vb.cell(r,c).value
                if p!=q and not(isinstance(p,(int,float)) and isinstance(q,(int,float)) and abs(p-q)<=1e-9*max(1,abs(p))):
                    d['value(cached)']+=1; ex.append((wa.cell(r,c).coordinate,'cached',str(p)[:40],str(q)[:40]))
    print(s,dict(d)); 
    for e in ex[:int(sys.argv[3]) if len(sys.argv)>3 else 6]: print('   ',e)
