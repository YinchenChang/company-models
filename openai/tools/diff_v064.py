#!/usr/bin/env python3
"""v0.6.4 逐格比較（表 2：範圍外變動）。以穩定 ID 或列位對齊 v0.6.3 與 v0.6.4，依工作單第 2 節把每個變動格分成『預期』與『範圍外』。

用法：python3 tools/diff_v064.py <v0.6.3.xlsx> <v0.6.4.xlsx> [--md out.md]
範圍外＝不屬於工作單第 2 節『預期變動』的格；應為 0。
"""
import argparse
import re
import sys
from collections import Counter, OrderedDict

import openpyxl

ROW_PAT = re.compile(r"(?<=[A-Z\$])\d+")


def norm(f):
    return ROW_PAT.sub("#", f) if isinstance(f, str) and f.startswith("=") else f


def same(p, q):
    if isinstance(p, (int, float)) and isinstance(q, (int, float)) and not isinstance(p, bool):
        return abs(p - q) <= 1e-9 * max(1, abs(p), abs(q))
    return p == q


def load(path):
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("old")
    ap.add_argument("new")
    ap.add_argument("--md")
    a = ap.parse_args()
    (fo, vo), (fn, vn) = load(a.old), load(a.new)
    rep = OrderedDict()
    out_scope = []

    def rec(sheet, kind, n=1):
        rep.setdefault(sheet, Counter())[kind] += n

    # 1. 完全不應變動的頁（逐格：公式文字＋重算值）
    for s in ("Derived_V9", "Demand", "Revenue", "Compute", "Cost", "Funding", "Reverse", "Map_v05"):
        wo, wn, ao, an = fo[s], fn[s], vo[s], vn[s]
        for r in range(1, max(wo.max_row, wn.max_row) + 1):
            for c in range(1, max(wo.max_column, wn.max_column) + 1):
                x, y = wo.cell(r, c).value, wn.cell(r, c).value
                p, q = ao.cell(r, c).value, an.cell(r, c).value
                if x != y or not same(p, q):
                    rec(s, "範圍外")
                    out_scope.append((s, wo.cell(r, c).coordinate, x, y, p, q))
                else:
                    rec(s, "相同")
    # 2. SRC_OAI／Inputs：以 ID 對齊（欄 A）
    allowed_inp = {"INP_289", "INP_290", "INP_291", "INP_303"}           # W1 改值、W3 改註記
    for s, ids_new in (("SRC_OAI", None), ("Inputs", None)):
        wo, wn, ao, an = fo[s], fn[s], vo[s], vn[s]
        old = {wo.cell(r, 1).value: r for r in range(5, wo.max_row + 1) if wo.cell(r, 1).value}
        new = {wn.cell(r, 1).value: r for r in range(5, wn.max_row + 1) if wn.cell(r, 1).value}
        for k in old:
            if k not in new:
                rec(s, "預期：列移除（W4）" if k in ("SRC_OAI_141", "SRC_OAI_142", "SRC_OAI_143") else "範圍外：列消失")
                if k not in ("SRC_OAI_141", "SRC_OAI_142", "SRC_OAI_143"):
                    out_scope.append((s, k, "列消失", "", "", ""))
        for k in new:
            if k not in old:
                rec(s, "預期：新增列")
                continue
            ro, rn = old[k], new[k]
            for c in range(1, wo.max_column + 1):
                x, y = norm(wo.cell(ro, c).value), norm(wn.cell(rn, c).value)
                p, q = ao.cell(ro, c).value, an.cell(rn, c).value
                if x != y or not same(p, q):
                    if s == "Inputs" and k in allowed_inp:
                        rec(s, "預期：" + k)
                    else:
                        rec(s, "範圍外")
                        out_scope.append((s, f"{k} col{c}", x, y, p, q))
                else:
                    rec(s, "相同")
    # 3. TK_Link：版本字串可變、數值不可變
    wo, wn, ao, an = fo["TK_Link"], fn["TK_Link"], vo["TK_Link"], vn["TK_Link"]
    ver_cells = {"B3", "B4", "B5", "B6", "B7"}
    for r in range(1, max(wo.max_row, wn.max_row) + 1):
        for c in range(1, max(wo.max_column, wn.max_column) + 1):
            x, y = wo.cell(r, c).value, wn.cell(r, c).value
            coord = wo.cell(r, c).coordinate
            if x == y:
                rec("TK_Link", "相同")
            elif coord in ver_cells or (r >= 10 and c in (7, 8, 9)):
                rec("TK_Link", "預期：版本／SHA／讀取日字串")
            else:
                rec("TK_Link", "範圍外")
                out_scope.append(("TK_Link", coord, x, y, "", ""))
    # 4. PxQ：列 1–298 位置不變；變動欄須在預期集合內
    wo, wn, ao, an = fo["PxQ"], fn["PxQ"], vo["PxQ"], vn["PxQ"]
    scope1 = set("FGH")
    exp_rows_any = {1, 9, 15, 26, 253, 283, 287, 297, 34, 35, 213}          # 標題、頂部口徑一、讀法、X205、X228、X232、X240、GW 列、營收列
    for r in range(1, 299):
        for c in range(1, 17):
            x, y = wo.cell(r, c).value, wn.cell(r, c).value
            p, q = ao.cell(r, c).value, an.cell(r, c).value
            if x == y and same(p, q):
                rec("PxQ", "相同")
                continue
            col = wo.cell(r, c).column_letter
            if col in scope1 or r in exp_rows_any or col in ("I", "J", "K") and r == 35 or col == "P" or col in "DE" and r in (9, 15):
                rec("PxQ", "預期")
            elif col == "H" and r == 253:
                rec("PxQ", "預期")
            else:
                rec("PxQ", "範圍外")
                out_scope.append(("PxQ", wo.cell(r, c).coordinate, x, y, p, q))
    rec("PxQ", "預期：新增列（第 8.4 節、第九節）", wn.max_row - 298)
    # 5. Checks：以編號對齊
    wo, wn, ao, an = fo["Checks"], fn["Checks"], vo["Checks"], vn["Checks"]
    old = {wo.cell(r, 1).value: r for r in range(5, wo.max_row + 1) if wo.cell(r, 1).value and str(wo.cell(r, 1).value).startswith("C")}
    new = {wn.cell(r, 1).value: r for r in range(5, wn.max_row + 1) if wn.cell(r, 1).value and str(wn.cell(r, 1).value).startswith("C")}
    chk_expected = {"C05", "C129", "C130", "C160", "C164", "C172", "C175"}      # 列數、TK 版本/SHA 顯示、版本字串、口徑一比值與倍數顯示
    for k in old:
        if k not in new:
            rec("Checks", "範圍外：編號消失")
            out_scope.append(("Checks", k, "消失", "", "", ""))
            continue
        diff = False
        for c in range(1, 7):
            x, y = norm(wo.cell(old[k], c).value), norm(wn.cell(new[k], c).value)
            p, q = ao.cell(old[k], c).value, an.cell(new[k], c).value
            if x != y or not same(p, q):
                diff = True
                if k not in chk_expected:
                    out_scope.append(("Checks", f"{k} col{c}", x, y, p, q))
        rec("Checks", ("預期" if diff else "相同"))
    for k in new:
        if k not in old:
            rec("Checks", "預期：新增檢查")
    # 6. 具名範圍
    dn_old, dn_new = set(fo.defined_names), set(fn.defined_names)
    rec("具名範圍", f"新增 {len(dn_new - dn_old)}", 0)
    lines = ["| 頁 | 統計 |", "|---|---|"] + [f"| {s} | {dict(c)} |" for s, c in rep.items()]
    lines.append(f"| 具名範圍 | 新增 {len(dn_new - dn_old)}：{sorted(dn_new - dn_old)}；移除 {len(dn_old - dn_new)}：{sorted(dn_old - dn_new)} |")
    n_out = len(out_scope)
    lines.append(f"\n範圍外變動：{n_out}")
    for t in out_scope[:40]:
        lines.append("- " + " | ".join(str(z)[:80] for z in t))
    text = "\n".join(lines)
    print(text)
    if a.md:
        open(a.md, "w", encoding="utf-8").write(text + "\n")
    sys.exit(0)


if __name__ == "__main__":
    main()
