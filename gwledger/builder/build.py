# GW 帳（全市場 GW Ledger）建檔程式
# 用法：
#   首版（由種子建）：python3 gwledger/builder/build.py --out build/raw/<YYYYMMDD>_GWLedger_v0.1.xlsx
#   之後各版（Excel 優先）：python3 gwledger/builder/build.py --base gwledger/model/<現行>.xlsx --out <raw.xlsx>
#     --base 時，In 頁的藍字輸入（口徑、標記、來源、說明、各年值或低／基準／高）一律讀回現行 Excel；種子只用於 Excel 裡還沒有的鍵。
# 產出後以 LibreOffice headless 重算存檔（見 gwledger/README.md），重算錯誤須為 0。
# 原則（比照 siliconomics/builder/build.py）：Excel 是唯一事實來源；本程式只產生公式頁與格式，計算一律為 Excel 公式。
import argparse, json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
GL = os.path.dirname(HERE)
ROOT = os.path.dirname(GL)
VERSION = "v0.1"
DATE = "20261028"
TK_SNAPSHOT = os.path.join(ROOT, "siliconomics", "data", "tokenomics_snapshot_v21.7.json")
SIL_DIR = os.path.join(ROOT, "siliconomics", "data")
SEED = os.path.join(HERE, "seed", "seed_v0.1.json")

BLUE = Font(color="0000FF"); GREEN = Font(color="008000"); BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7"); KEY = PatternFill("solid", fgColor="FFF2CC")
PEND = PatternFill("solid", fgColor="FFFF00"); WRAP = Alignment(wrap_text=True, vertical="top")
YEARS = list(range(2023, 2033))           # 2023＝期初存量年；輸出 2024–2032
OUTY = [y for y in YEARS if y >= 2024]
YC = {y: get_column_letter(4 + i) for i, y in enumerate(YEARS)}   # 計算頁：D..M
IC = {y: get_column_letter(8 + i) for i, y in enumerate(YEARS)}   # In 頁：H..Q
SCEN = ["A", "B", "C"]
DLY = ["D0", "D1", "D2"]
GENS = ["H100", "GB200", "GB300", "VR200", "RU"]
GEN_NAME = {"H100": "Hopper H100", "GB200": "GB200 NVL72", "GB300": "GB300 NVL72", "VR200": "VR200 NVL72", "RU": "Rubin Ultra"}


def hdr(ws, row, labels, c0=1):
    for j, h in enumerate(labels, c0):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def load_base_inputs(path):
    """讀回現行 Excel 的 In 頁（Excel 優先）：{鍵: {欄: 值}}（C、E–Q 欄）"""
    if not path:
        return {}
    ws = openpyxl.load_workbook(path, data_only=False)["In"]
    out = {}
    for r in range(5, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if not k or not ws.cell(r, 19).value:
            continue
        out[k] = {c: ws.cell(r, c).value for c in [3] + list(range(5, 18))}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base")
    ap.add_argument("--out", required=True)
    ap.add_argument("--version")
    ap.add_argument("--date")
    a = ap.parse_args()
    if a.base:
        fn = os.path.basename(a.base)
        if a.version is None and "_v" in fn:
            a.version = "v" + fn.rsplit("_v", 1)[1].rsplit(".xlsx", 1)[0]
        if a.date is None and fn[:8].isdigit():
            a.date = fn[:8]
    a.version = a.version or VERSION
    a.date = a.date or DATE
    seed = json.load(open(SEED, encoding="utf-8"))
    base = load_base_inputs(a.base)
    tk = json.load(open(TK_SNAPSHOT, encoding="utf-8"))
    sil_fn = open(os.path.join(SIL_DIR, "CURRENT_SNAPSHOT"), encoding="utf-8").read().strip()
    sil = json.load(open(os.path.join(SIL_DIR, sil_fn), encoding="utf-8"))
    P = seed["players"]; PK = [p["key"] for p in P]; PN = {p["key"]: p["name"] for p in P}
    U = seed["users"]

    wb = openpyxl.Workbook()
    RD = wb.active; RD.title = "README"
    IN = wb.create_sheet("In")
    TKS = wb.create_sheet("TK_In")
    SL = wb.create_sheet("SIL_In")
    PL = wb.create_sheet("Players")
    MX = wb.create_sheet("Matrix")
    PP = wb.create_sheet("Pipeline")
    DP = wb.create_sheet("Deploy")
    C0 = wb.create_sheet("C0")
    C1 = wb.create_sheet("C1")
    CP = wb.create_sheet("Compare")
    IF = wb.create_sheet("Interface")
    CK = wb.create_sheet("Checks")
    DC = wb.create_sheet("Decisions")
    EV = wb.create_sheet("DB_Evidence")
    SR = wb.create_sheet("Sources")
    RQ = wb.create_sheet("Requests")

    # ================= In：藍字輸入（Excel 擁有） =================
    IN["A1"] = ("In — GW 帳輸入（藍字＝輸入，由本活頁簿擁有；要改值直接改本頁，builder 重建時以 --base 讀回）。"
                "純量列：H／I／J＝低／基準／高，具名範圍＝鍵（I 欄）；年度列：H–Q＝2023–2032；分攤表：H–M＝使用者；園區表見節標題")
    IN["A2"] = ("標記：Verified／Interested-party（寫明是誰）／Analogy／Assumed（給低／高）／Derived（寫明換算）／待證據（值留空）。"
                "來源 ID 見 Sources 頁。口徑：IT／設施／未說明（未說明者基準取 IT 讀法，Players 另列設施讀法）。S 欄＝輸入列類型（治理檢查用）")
    hdr(IN, 4, ["鍵", "項目", "口徑", "單位", "標記", "來源 ID", "說明"] + [str(y) for y in YEARS] + ["", "類型"])
    INROW = {}; INKIND = {}
    r = 5

    def B(key, col, default):
        if key in base and col in base[key]:
            return base[key][col]
        return default

    def sec(title):
        nonlocal r
        IN.cell(r, 1, title).font = BOLD; r += 1

    def meta(key, label, basis, unit, tag, src, note, kind):
        IN.cell(r, 1, key); IN.cell(r, 2, label)
        IN.cell(r, 3, B(key, 3, basis)); IN.cell(r, 4, unit)
        IN.cell(r, 5, B(key, 5, tag)); IN.cell(r, 6, B(key, 6, src)); IN.cell(r, 7, B(key, 7, note))
        IN.cell(r, 19, kind)
        if B(key, 5, tag) == "待證據":
            for j in range(1, 8): IN.cell(r, j).fill = PEND
        INROW[key] = r; INKIND[key] = kind

    def scalar(s):
        nonlocal r
        k = s["key"]
        meta(k, s["label"], s.get("basis", "—"), s["unit"], s["tag"], s["src"], s.get("note", ""), 1)
        for col, v in ((8, s.get("lo")), (9, s.get("base")), (10, s.get("hi"))):
            v = B(k, col, v)
            if v is not None:
                c = IN.cell(r, col, v); c.font = BLUE
                if s["unit"] == "%": c.number_format = "0.0%"
        name(wb, k, f"In!$I${r}")
        if B(k, 8, s.get("lo")) is not None: name(wb, k + "_lo", f"In!$H${r}")
        if B(k, 10, s.get("hi")) is not None: name(wb, k + "_hi", f"In!$J${r}")
        r += 1

    def yrow(key, label, basis, unit, tag, src, note, vals, kind=2, fmt="#,##0.0##"):
        nonlocal r
        meta(key, label, basis, unit, tag, src, note, kind)
        for y in YEARS:
            v = B(key, 8 + YEARS.index(y), vals.get(str(y)))
            if v is not None:
                c = IN.cell(r, 8 + YEARS.index(y), v); c.font = BLUE; c.number_format = fmt
        r += 1

    for s in seed["scalars"]:
        if "sec" in s:
            sec(s["sec"]); continue
        scalar(s)
    sec("S5. 玩家非 NVIDIA 占比覆寫（未列者＝市場 NonNV 比例）")
    for p, o in seed["nnv"].items():
        if p.startswith("_"): continue
        scalar(dict(key=f"nnv_{p}", label=f"{PN[p]} 新增 GW 非 NVIDIA 占比", unit="%", base=o["base"], tag=o["tag"], src=o["src"], note=o["note"]))
    sec("S6. 個股模型財年結束月（C1 曆年→財年換算）")
    FYE = {}
    for p, o in seed["capex"].items():
        if p.startswith("_"): continue
        scalar(dict(key=f"fye_{p}", label=f"{PN[p]} 財年結束月", unit="月", base=o["fye"], tag="Verified", src=o["src"], note="company.json calendar.fiscalYearEndMonth"))
        FYE[p] = f"fye_{p}"
    S_LAST = r - 1

    sec("Y1. 玩家年底上線 MW（營運者口徑；2023–2026 實際或最新，2027–2032 公告路徑；空白＝未知，Players 頁依 GL01-03 補值）")
    for p in PK:
        o = seed["mw"].get(p)
        if o:
            yrow(f"mw_{p}", f"{PN[p]} 年底上線 MW", o["basis"], "MW", o["tag"], o["src"], o["note"], o["y"], fmt="#,##0.0")
        else:
            yrow(f"mw_{p}", f"{PN[p]} 年底上線 MW", "未說明", "MW", "待證據", "SRC_GL_000", "待證據（evidence_raw 複核後填入）", {}, fmt="#,##0.0")
    sec("Y2. 個股模型資本支出（公司財年；欄＝財年，US$B；PR #190 未合併前讀現行 company.json／data）")
    for p in PK:
        o = seed["capex"].get(p)
        if o and o["y"]:
            yrow(f"cap_{p}", f"{PN[p]} 資本支出（財年）", "—", "US$B", o["tag"], o["src"], o["note"], o["y"], fmt="#,##0.0")
        else:
            yrow(f"cap_{p}", f"{PN[p]} 資本支出（財年）", "—", "US$B", "待證據", "SRC_GL_000", (o or {}).get("note", "待證據（無個股模型或未取得）"), {}, fmt="#,##0.0")
    sec("Y3. 非 AI 資本支出（公司財年，US$B；C1 加回）")
    for p in PK:
        yrow(f"nai_{p}", f"{PN[p]} 非 AI 資本支出（財年）", "—", "US$B", "待證據", "SRC_GL_000", "待證據；空白視為 0（C1 推算偏低）", {}, fmt="#,##0.0")
    sec("Y4. 使用者總量（報導或使用者模型；只作 Matrix 欄和對照）")
    for o in seed["usertotals"]:
        yrow(o["key"], o["label"], o["basis"], "GW", o["tag"], o["src"], o["note"], o["y"], fmt="0.00")
    sec("Y5. 第三方總量（C0 佐證；只作佐證，不校準）")
    for o in seed["thirdparty"]:
        yrow(o["key"], o["label"], o["basis"], "GW", "待證據", "SRC_GL_000", "待證據（B_labs_china_totals 複核後填入）", {}, fmt="0.00")
    sec("Y6. 世代組合（新增 GW；NVIDIA 內部占比；空白年＝Players 頁由 SIL_In 推導〔Derived〕）")
    mx = seed["mix"]
    for g in GENS:
        yrow(f"mix_{g}", f"新增 GW 世代占比：{GEN_NAME[g]}", "—", "%", mx["tag_nv"], mx["src_nv"], mx["note_nv"], mx["nv"].get(g, {}), fmt="0%")
    yrow("nnv_mkt", "市場新增 GW 非 NVIDIA 占比（覆寫；空白＝SIL 推導，2023–25 沿用 2026、2030–32 沿用 2029）", "—", "%", "待證據", "SRC_GL_000",
         "2026–29 預設由 Siliconomics Units × 每顆 IT capex 推導（Derived）；2023–25 無來源，沿用 2026 值待決", {}, fmt="0%")
    Y_LAST = r - 1

    sec("T1. 營運者 → 使用者分攤（列＝營運者；H–M＝" + "、".join(u["name"] for u in U) + "；空白＝0；1 − 合計＝未分配）")
    hdr(IN, r, ["鍵", "營運者", "口徑", "單位", "標記", "來源 ID", "說明"] + [u["name"] for u in U]); r += 1
    for p in PK:
        o = seed["shares"].get(p)
        k = f"shr_{p}"
        if o:
            meta(k, f"{PN[p]} 使用者分攤", "—", "比", o["tag"], o["src"], o["note"] + (f"（自用區間 {o['lo']}–{o['hi']}）" if "lo" in o else ""), 3)
        else:
            meta(k, f"{PN[p]} 使用者分攤", "—", "比", "待證據", "SRC_GL_000", "待證據", 3)
        for j, u in enumerate(U):
            v = B(k, 8 + j, (o or {}).get("u", {}).get(u["key"]))
            if v is not None:
                c = IN.cell(r, 8 + j, v); c.font = BLUE; c.number_format = "0%"
        r += 1
    SH_LAST = r - 1
    r += 1
    sec("T2. 園區管線（H 營運者鍵、I 使用者鍵、J MW、K 公告上線年、L 公告上線月、M 已知延遲〔月〕、N 納入〔1＝計入 Pipeline〕）")
    hdr(IN, r, ["鍵", "園區", "口徑", "單位", "標記", "來源 ID", "說明", "營運者", "使用者", "MW", "上線年", "上線月", "已知延遲月", "納入"]); r += 1
    CAMP = []
    for cp in seed["campus"]:
        k = cp["key"]
        meta(k, cp["name"], cp["basis"], "MW", cp["tag"], cp["src"], cp["note"], 4)
        for j, f in enumerate(("op", "user", "mw", "year", "month", "delay", "inc")):
            v = B(k, 8 + j, cp.get(f))
            if v is not None:
                c = IN.cell(r, 8 + j, v); c.font = BLUE
        CAMP.append((k, r)); r += 1
    IN_LAST = r - 1
    for col, w in zip("ABCDEFGHIJKLMNOPQRS", (16, 50, 8, 7, 18, 12, 70) + (9,) * 11 + (5,)):
        IN.column_dimensions[col].width = w

    # ================= TK_In：Tokenomics 鏡像（只讀） =================
    src = tk["source"]
    TKS["A1"] = (f"TK_In — Tokenomics 取數（只讀鏡像；權威在 Tokenomics）：{src['file']}（{src['version']}，commit {src['commit'][:7]}，擷取 {src['extractedAt']}）。"
                 "同一份快照與 Siliconomics 相同（siliconomics/data/tokenomics_snapshot_v21.7.json），不另抽；綠字不得修改")
    TKS["A2"] = "快照註記：" + src.get("note", "")
    hdr(TKS, 4, ["名稱（TK）", "世代／項目", "本庫名稱", "低", "基準", "高", "單位", "TK 位置"])
    gens = {g["name"]: g["code"] for g in src["generations"]}
    rr = 5; TKN = set()
    for nm_, it in tk["items"].items():
        short = nm_[3:]
        if it.get("kind") == "gen_cost":
            for g, vv in it["values"].items():
                sname = f"TK_{short}_{gens.get(g, g)}"
                TKS.cell(rr, 1, nm_); TKS.cell(rr, 2, g); TKS.cell(rr, 3, sname)
                for j, c_ in ((4, "低成本"), (5, "基準"), (6, "高成本")):
                    TKS.cell(rr, j, vv[c_]).font = GREEN
                TKS.cell(rr, 7, it.get("unit")); TKS.cell(rr, 8, it.get("ref"))
                name(wb, sname, f"TK_In!$E${rr}"); name(wb, sname + "_lo", f"TK_In!$D${rr}"); name(wb, sname + "_hi", f"TK_In!$F${rr}")
                TKN.add(sname); rr += 1
        else:
            vals = it.get("values") or it.get("value")
            sname = f"TK_{short}"
            TKS.cell(rr, 1, nm_); TKS.cell(rr, 2, it.get("label")); TKS.cell(rr, 3, sname)
            if isinstance(vals, list):
                for j, x in zip((4, 5, 6), vals): TKS.cell(rr, j, x).font = GREEN
            else:
                TKS.cell(rr, 5, vals).font = GREEN
            TKS.cell(rr, 7, it.get("unit")); TKS.cell(rr, 8, it.get("ref"))
            name(wb, sname, f"TK_In!$E${rr}"); name(wb, sname + "_lo", f"TK_In!$D${rr}"); name(wb, sname + "_hi", f"TK_In!$F${rr}")
            TKN.add(sname); rr += 1
    TK_LAST = rr - 1
    HAS_FACGW = "IF_FacilityGW" in tk["items"]
    for col, w in zip("ABCDEFGH", (22, 40, 26, 10, 10, 10, 12, 20)):
        TKS.column_dimensions[col].width = w

    # ================= SIL_In：Siliconomics 鏡像（只讀） =================
    sn = sil["names"]; ssrc = sil["source"]
    SL["A1"] = (f"SIL_In — Siliconomics 取數（只讀鏡像；權威在 AI半導體）：{ssrc['file']}（{ssrc['version']}，sha256 {ssrc['sha256'][:12]}）。"
                "由 siliconomics/data/CURRENT_SNAPSHOT 指向；綠字不得修改")
    hdr(SL, 3, ["名稱（SIF_）", "本庫名稱", "單位", "值（純量）／2023"] + [str(y) for y in YEARS[1:]])
    SILS = {}; SILV = {}
    rr = 4
    SL.cell(rr, 1, "A. 純量").font = BOLD; rr += 1
    scal = [("SIF_vr26", "x"), ("SIF_vr26_lo", "x"), ("SIF_vr26_hi", "x")] + [(f"SIF_it_{k}", "US$K／顆") for k in ("GB", "VR", "RU", "HP", "AMD8", "AMDR", "TPU", "TRN", "ASIC")]
    for nm_, unit in scal:
        nn = "SIL_" + nm_[4:]
        SL.cell(rr, 1, nm_); SL.cell(rr, 2, nn); SL.cell(rr, 3, unit)
        SL.cell(rr, 4, sn[nm_]).font = GREEN
        name(wb, nn, f"SIL_In!$D${rr}"); SILS[nm_] = nn; rr += 1
    rr += 1
    SL.cell(rr, 1, "B. 年度（欄 D–M＝2023–2032；Siliconomics 只到 2029）").font = BOLD; rr += 1
    vec = []
    for k in ("bw", "rb", "ru", "hp", "amd", "tpu", "trn", "oth"):
        vec.append((f"SIF_UnitsPlat_{k}", 2025, "百萬顆"))
    for s in SCEN:
        vec.append((f"SIF_UnitsCoef_{s}", 2025, "x"))
    for s in SCEN:
        vec.append((f"SIF_DeployGW_{s}", 2026, "GW（成本等值）"))
    for s in SCEN:
        vec.append((f"SIF_AIITCapex_{s}", 2026, "US$B"))
    vec.append(("SIF_ITCapexPerGW", 2026, "US$B/GW"))
    for nm_, y0, unit in vec:
        SL.cell(rr, 1, nm_); SL.cell(rr, 3, unit)
        for i, v in enumerate(sn[nm_]):
            y = y0 + i
            if v is not None:
                SL.cell(rr, YEARS.index(y) + 4, v).font = GREEN
                SILV[(nm_, y)] = f"SIL_In!${YC[y]}${rr}"
        rr += 1
    for col, w in zip("ABCDEFGHIJKLM", (26, 16, 14, 14) + (10,) * 9):
        SL.column_dimensions[col].width = w

    def sv(nm_, y):
        return SILV[(nm_, y)]

    # ================= Players：玩家 × 年 =================
    PL["A1"] = ("Players — 玩家 × 年（2023–2032）年底上線 TK GW（甲：公告路徑；營運者口徑）。計算 MW：有輸入取輸入；無輸入而最新觀測在該年取觀測；"
                "兩已知值間線性內插；最後已知值後持平沿用；第一個已知值前留空（GL01-03）。TK GW＝設施電力＝IT MW ÷ 1000 ÷ g（g＝每 TK GW 含 IT GW，依世代組合加權）")
    PL["A2"] = "旗標：1 輸入／2 觀測／3 內插／4 沿用／0 未知。年新增＝年底差（第一個已知年為『新納入存量』，不計新增）"
    hdr(PL, 3, ["項目", "鍵", "單位"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    PL.cell(r, 1, "世代組合（新增 GW；NVIDIA 內部占比）與換算係數").font = BOLD; r += 1
    MIXR = {}
    RUSH = r + len(GENS) + 1
    for g in GENS:
        ir = INROW[f"mix_{g}"]
        PL.cell(r, 1, f"世代占比：{GEN_NAME[g]}"); PL.cell(r, 2, f"mix_{g}"); PL.cell(r, 3, "%")
        for y in YEARS:
            c = YC[y]; inref = f"In!${IC[y]}${ir}"
            if y <= 2025:
                d = "0"
            elif y == 2026:
                d = {"GB300": "(1-SIL_vr26)", "VR200": "SIL_vr26"}.get(g, "0")
            elif y <= 2029:
                d = {"VR200": f"(1-{c}${RUSH})", "RU": f"{c}${RUSH}"}.get(g, "0")
            else:
                d = f"{YC[2029]}{r}"
            PL.cell(r, YEARS.index(y) + 4, f'=IF({inref}<>"",{inref},{d})').number_format = "0%"
        MIXR[g] = r; r += 1
    PL.cell(r, 1, "世代占比合計（須＝1）"); PL.cell(r, 3, "%")
    for y in YEARS:
        PL.cell(r, YEARS.index(y) + 4, f"=SUM({YC[y]}{MIXR[GENS[0]]}:{YC[y]}{MIXR[GENS[-1]]})").number_format = "0%"
    MIXSUM = r; r += 1
    assert r == RUSH
    PL.cell(r, 1, "Rubin Ultra 占 Rubin 平台 GW 比（Derived：Units ÷ TK 每 GW GPU 數）"); PL.cell(r, 3, "%")
    for y in (2027, 2028, 2029):
        ru, rb = sv("SIF_UnitsPlat_ru", y), sv("SIF_UnitsPlat_rb", y)
        PL.cell(r, YEARS.index(y) + 4, f"=({ru}/TK_GPUsPerGW_RU)/({rb}/TK_GPUsPerGW_VR200+{ru}/TK_GPUsPerGW_RU)").number_format = "0.0%"
    PL.cell(r, 14, "工作單：Rubin Ultra 比例讀 Siliconomics Units；以每 GW GPU 數換成 GW 占比（GL01-11）；2027 Blackwell 2.4M 顆不計（工作單：2027 起 VR200）")
    r += 1
    NNVR = r
    PL.cell(r, 1, "市場新增 GW 非 NVIDIA 占比"); PL.cell(r, 2, "nnv_mkt"); PL.cell(r, 3, "%")
    ir = INROW["nnv_mkt"]
    for y in YEARS:
        inref = f"In!${IC[y]}${ir}"
        if 2026 <= y <= 2029:
            amd = "SIL_it_AMD8" if y == 2026 else "SIL_it_AMDR"
            nnv = f"({sv('SIF_UnitsPlat_amd', y)}*{amd}+{sv('SIF_UnitsPlat_tpu', y)}*SIL_it_TPU+{sv('SIF_UnitsPlat_trn', y)}*SIL_it_TRN+{sv('SIF_UnitsPlat_oth', y)}*SIL_it_ASIC)"
            nv = f"({sv('SIF_UnitsPlat_bw', y)}*SIL_it_GB+{sv('SIF_UnitsPlat_rb', y)}*SIL_it_VR+{sv('SIF_UnitsPlat_ru', y)}*SIL_it_RU+{sv('SIF_UnitsPlat_hp', y)}*SIL_it_HP)"
            d = f"{nnv}/({nnv}+{nv})"
        elif y < 2026:
            d = f"{YC[2026]}{r}"
        else:
            d = f"{YC[2029]}{r}"
        PL.cell(r, YEARS.index(y) + 4, f'=IF({inref}<>"",{inref},{d})').number_format = "0.0%"
    PL.cell(r, 14, "Derived：Siliconomics Units 非 NVIDIA 平台 IT capex 占比（工作單第 2 節）；2023–25 沿用 2026（待決）、2030–32 沿用 2029")
    r += 1
    GROW = r
    PL.cell(r, 1, "g：每 TK GW（設施）含 IT GW（世代加權）"); PL.cell(r, 3, "x")
    for y in YEARS:
        c = YC[y]
        PL.cell(r, YEARS.index(y) + 4, f"=gIT_oth+(gIT_vr200-gIT_oth)*{c}{MIXR['VR200']}*(1-{c}{NNVR})").number_format = "0.0000"
    PL.cell(r, 14, "VR200＝gIT_vr200（規劃書 919.8 MW IT／TK GW）；其他世代與非 NVIDIA＝gIT_oth（v0.1 沿用同值，Assumed）")
    r += 2
    PR = {}
    OBS = {p for p in PK if f"obs_{p}" in INROW}
    for p in PK:
        ir = INROW[f"mw_{p}"]
        PL.cell(r, 1, PN[p]).font = BOLD; PL.cell(r, 2, p); PL.cell(r, 14, f"In 第 {ir} 列；口徑＝In!C{ir}")
        r += 1
        rows = {}
        for k, lab, unit in (("in", "輸入 MW", "MW"), ("k", "已知值（輸入或觀測）", "MW"), ("L", "前一已知值", "MW"), ("LY", "前一已知年", "年"),
                             ("N", "下一已知值", "MW"), ("NY", "下一已知年", "年"), ("v", "計算 MW（營運者口徑）", "MW"), ("f", "旗標", "—"),
                             ("tk", "年底上線 TK GW（設施）", "GW"), ("alt", "設施讀法（口徑未說明時不 ÷ g）", "GW"), ("avg", "年平均 TK GW", "GW"),
                             ("add", "年新增 TK GW", "GW"), ("stk", "新納入存量 TK GW", "GW"), ("nnv", "新增 GW 非 NVIDIA 占比", "%")):
            rows[k] = r; PL.cell(r, 1, "　" + lab); PL.cell(r, 2, f"{p}_{k}"); PL.cell(r, 3, unit); r += 1
        for i, y in enumerate(YEARS):
            c = YC[y]; inref = f"In!${IC[y]}${ir}"; col = i + 4
            pc = YC.get(y - 1); nc = YC.get(y + 1)
            PL.cell(rows["in"], col, f'=IF({inref}<>"",{inref},"")').font = GREEN
            if p in OBS:
                kf = f'=IF({inref}<>"",{inref},IF(AND(In!$I${INROW["obs_" + p]}<>"",INT(obsd_{p})={y}),obs_{p},""))'
            else:
                kf = f'=IF({inref}<>"",{inref},"")'
            PL.cell(rows["k"], col, kf)
            K = f"{c}{rows['k']}"
            if pc is None:
                PL.cell(rows["L"], col, f'=IF({K}<>"",{K},"")'); PL.cell(rows["LY"], col, f'=IF({K}<>"",{y},"")')
            else:
                PL.cell(rows["L"], col, f'=IF({K}<>"",{K},{pc}{rows["L"]})'); PL.cell(rows["LY"], col, f'=IF({K}<>"",{y},{pc}{rows["LY"]})')
            if nc is None:
                PL.cell(rows["N"], col, f'=IF({K}<>"",{K},"")'); PL.cell(rows["NY"], col, f'=IF({K}<>"",{y},"")')
            else:
                PL.cell(rows["N"], col, f'=IF({K}<>"",{K},{nc}{rows["N"]})'); PL.cell(rows["NY"], col, f'=IF({K}<>"",{y},{nc}{rows["NY"]})')
            Lr, LYr, Nr, NYr = (f"{c}{rows[x]}" for x in ("L", "LY", "N", "NY"))
            PL.cell(rows["v"], col, f'=IF({K}<>"",{K},IF(AND({Lr}<>"",{Nr}<>""),{Lr}+({Nr}-{Lr})*({y}-{LYr})/({NYr}-{LYr}),IF({Lr}<>"",{Lr},"")))').number_format = "#,##0"
            PL.cell(rows["f"], col, f'=IF({inref}<>"",1,IF({K}<>"",2,IF(AND({Lr}<>"",{Nr}<>""),3,IF({Lr}<>"",4,0))))')
            V = f"{c}{rows['v']}"
            cc = PL.cell(rows["tk"], col, f'=IF({V}="","",IF(In!$C${ir}="設施",{V}/1000,{V}/1000/{c}${GROW}))'); cc.number_format = "0.000"; cc.fill = KEY
            PL.cell(rows["alt"], col, f'=IF({V}="","",IF(In!$C${ir}="未說明",{V}/1000,{c}{rows["tk"]}))').number_format = "0.000"
            T = f"{c}{rows['tk']}"
            if pc is None:
                PL.cell(rows["avg"], col, '=""'); PL.cell(rows["add"], col, 0)
                PL.cell(rows["stk"], col, f'=IF({T}<>"",{T},0)').number_format = "0.000"
            else:
                PT = f"{pc}{rows['tk']}"
                PL.cell(rows["avg"], col, f'=IF(AND({PT}<>"",{T}<>""),({PT}+{T})/2,"")').number_format = "0.000"
                PL.cell(rows["add"], col, f'=IF(AND({PT}<>"",{T}<>""),{T}-{PT},0)').number_format = "0.000"
                PL.cell(rows["stk"], col, f'=IF(AND({PT}="",{T}<>""),{T},0)').number_format = "0.000"
            PL.cell(rows["nnv"], col, (f"=nnv_{p}" if f"nnv_{p}" in INROW else f"={c}${NNVR}")).number_format = "0.0%"
        PR[p] = rows; r += 1
    PL.cell(r, 1, "市場合計（Σ 營運者；主要玩家＋其他）").font = BOLD; r += 1
    TOT = {}
    for k, lab, key in (("tk", "年底上線 TK GW（設施）", "ye"), ("alt", "年底上線 GW（設施讀法：口徑未說明者不 ÷ g）", "alt"),
                        ("it", "年底上線 IT GW", "it"), ("avg", "年平均 TK GW", "avg"), ("add", "年新增 TK GW", "add"),
                        ("stk", "新納入存量 TK GW", "stk"), ("cnt", "有值玩家數", "cnt")):
        PL.cell(r, 1, lab); PL.cell(r, 2, f"mkt_{key}")
        for i, y in enumerate(YEARS):
            c = YC[y]
            if key == "it":
                f = f"={c}{TOT['ye']}*{c}${GROW}"
            elif key == "avg":
                f = '=""' if i == 0 else f"=({YC[y - 1]}{TOT['ye']}+{c}{TOT['ye']})/2"
            elif key == "cnt":
                f = "=COUNT(" + ",".join(f"{c}{PR[p]['tk']}" for p in PK) + ")"
            else:
                f = "=SUM(" + ",".join(f"{c}{PR[p][k]}" for p in PK) + ")"
            cc = PL.cell(r, i + 4, f); cc.number_format = "0" if key == "cnt" else "0.000"
            if key == "ye": cc.fill = KEY
        TOT[key] = r; r += 1
    PL_LAST = r
    PL.column_dimensions["A"].width = 52; PL.column_dimensions["B"].width = 12; PL.column_dimensions["N"].width = 60
    for y in YEARS: PL.column_dimensions[YC[y]].width = 9

    # ================= Matrix：營運者 × 使用者 =================
    MX["A1"] = ("Matrix — 營運者 × 使用者（年底 TK GW）。列＝營運者（Players 年底值 × In T1 分攤），欄＝使用者；『未分配』＝1 − 分攤合計。"
                "列和＝Players；欄和合計＝總量（處理 OpenAI、Anthropic 在雲端的重複計算：市場總量只按營運者加總一次）")
    r = 3; MXR = {}
    for y in (2025, 2026, 2027):
        MX.cell(r, 1, f"{y} 年底").font = BOLD; r += 1
        hdr(MX, r, ["營運者", "鍵"] + [u["name"] for u in U] + ["未分配", "列和", "Players", "差"]); r += 1
        r0 = r
        for p in PK:
            sr = INROW[f"shr_{p}"]
            MX.cell(r, 1, PN[p]); MX.cell(r, 2, p)
            MX.cell(r, 11, f"=N(Players!{YC[y]}{PR[p]['tk']})").font = GREEN
            for j in range(len(U)):
                MX.cell(r, 3 + j, f"=$K{r}*N(In!{get_column_letter(8 + j)}{sr})").number_format = "0.000"
            MX.cell(r, 9, f"=$K{r}*(1-SUM(In!H{sr}:M{sr}))").number_format = "0.000"
            MX.cell(r, 10, f"=SUM(C{r}:I{r})").number_format = "0.000"
            MX.cell(r, 12, f"=J{r}-K{r}").number_format = "0.000000"
            MX.cell(r, 11).number_format = "0.000"
            r += 1
        r1 = r - 1
        MX.cell(r, 1, "欄和（使用者合計）").font = BOLD
        for j in range(3, 13):
            col = get_column_letter(j)
            MX.cell(r, j, f"=SUM({col}{r0}:{col}{r1})").number_format = "0.000"
        MX.cell(r, 12, f"=SUM(C{r}:I{r})-K{r}").number_format = "0.000000"
        MXR[y] = dict(r0=r0, r1=r1, sum=r); r += 1
        for uk, ukey, ucol in (("oai", "usr_oai", 4), ("ant", "usr_ant", 5)):
            ir = INROW[ukey]
            MX.cell(r, 1, f"對照：{ukey} 報導值（IT 讀法 ÷ g；口徑 In!C{ir}）")
            MX.cell(r, ucol, f'=IF(In!{IC[y]}{ir}="","",IF(In!$C${ir}="設施",In!{IC[y]}{ir},In!{IC[y]}{ir}/Players!{YC[y]}${GROW}))').number_format = "0.000"
            MX.cell(r, 10, f'=IF(OR({get_column_letter(ucol)}{r}="",{get_column_letter(ucol)}{r}=0),"",{get_column_letter(ucol)}{MXR[y]["sum"]}/{get_column_letter(ucol)}{r})').number_format = "0.00"
            MX.cell(r, 11, "← 矩陣欄和 ÷ 報導（< 1＝營運者分攤或玩家覆蓋不足）")
            r += 1
        r += 1
    MX_LAST = r
    MX.column_dimensions["A"].width = 46
    for col in "CDEFGHIJKL": MX.column_dimensions[col].width = 10

    # ================= Pipeline：園區管線 × 延遲情景 =================
    PP["A1"] = ("Pipeline — 上線事件 × 延遲情景。事件＝(1) 園區（In T2，有上線年月且納入＝1）＋(2) 各玩家各年『未分園區殘差』＝Players 年新增 − 該年已列園區；"
                "＋新納入存量（不延遲）。時間以小數年表示（年 Y 月 M 底＝Y＋M/12；Y 年底＝Y＋1）。延遲只作用於評價時點 asof 之後：D0＝0；D1＝MAX(已知延遲, dly_d1_m)；"
                "D2＝D1 之上綁定比例 dly_d2_bind 再延 dly_d2_m。殘差在年內均勻上線（觀測年在觀測時點前已上線，不延遲）")
    r = 3
    PP.cell(r, 1, "A. 園區（In T2 鏡像與上線時點）").font = BOLD; r += 1
    hdr(PP, r, ["園區鍵", "名稱", "營運者", "使用者", "MW", "口徑", "納入", "上線年", "上線月", "已知延遲月", "t 公告", "TK GW", "t D0", "t D1", "t D2 綁定", "D0 上線年", "有效納入"]); r += 1
    CR0 = r
    for k, ir in CAMP:
        PP.cell(r, 1, k); PP.cell(r, 2, f"=In!B{ir}").font = GREEN
        for j, col in ((3, "H"), (4, "I"), (5, "J"), (7, "N"), (8, "K"), (9, "L"), (10, "M")):
            PP.cell(r, j, f'=IF(In!{col}{ir}="","",In!{col}{ir})').font = GREEN
        PP.cell(r, 6, f"=In!C{ir}").font = GREEN
        PP.cell(r, 11, f'=IF(OR(H{r}="",I{r}=""),9999,H{r}+I{r}/12)').number_format = "0.000"
        PP.cell(r, 12, f'=IF(E{r}="",0,IF(F{r}="設施",E{r}/1000,E{r}/1000/gIT_oth))').number_format = "0.000"
        PP.cell(r, 13, f"=K{r}").number_format = "0.000"
        PP.cell(r, 14, f"=IF(K{r}>asof,K{r}+MAX(N(J{r}),dly_d1_m)/12,K{r})").number_format = "0.000"
        PP.cell(r, 15, f"=IF(K{r}>asof,N{r}+dly_d2_m/12,K{r})").number_format = "0.000"
        PP.cell(r, 16, f'=IF(K{r}=9999,"",ROUNDUP(K{r},0)-1)')
        PP.cell(r, 17, f"=IF(AND(N(G{r})=1,K{r}<9999,L{r}>0),1,0)")
        r += 1
    CR1 = r - 1
    rng = lambda col: f"${col}${CR0}:${col}${CR1}"
    r += 1
    PP.cell(r, 1, "B. 未分園區殘差（TK GW；Players 年新增 − 該營運者該年 D0 上線的已列園區）").font = BOLD; r += 1
    hdr(PP, r, ["玩家", "鍵", "類別"] + [str(y) for y in YEARS]); r += 1
    RES = {}; ROBS = {}; RUNI = {}
    for kind, store in (("殘差合計", RES), ("觀測型（不延遲）", ROBS), ("均勻型（延遲適用）", RUNI)):
        for p in PK:
            PP.cell(r, 1, PN[p]); PP.cell(r, 2, p); PP.cell(r, 3, kind)
            for y in OUTY:
                c = YC[y]
                if store is RES:
                    f = f'=Players!{c}{PR[p]["add"]}-SUMIFS({rng("L")},{rng("C")},"{p}",{rng("Q")},1,{rng("P")},{y})'
                elif store is ROBS:
                    f = f"=IF(Players!{c}{PR[p]['f']}=2,{c}{RES[p]},0)"
                else:
                    f = f"={c}{RES[p]}-{c}{ROBS[p]}"
                PP.cell(r, YEARS.index(y) + 4, f).number_format = "0.000"
            store[p] = r; r += 1
        r += 1
    AGG = {}
    for key, lab, f_ in (("res", "殘差合計", lambda c: f"=SUM({c}{RES[PK[0]]}:{c}{RES[PK[-1]]})"),
                         ("robs", "觀測型殘差合計", lambda c: f"=SUM({c}{ROBS[PK[0]]}:{c}{ROBS[PK[-1]]})"),
                         ("runi", "均勻型殘差合計", lambda c: f"=SUM({c}{RUNI[PK[0]]}:{c}{RUNI[PK[-1]]})"),
                         ("camp", "園區（D0 上線年）合計", None),
                         ("stk", "新納入存量（Players）", lambda c: f"=Players!{c}{TOT['stk']}"),
                         ("addchk", "殘差＋園區（須＝Players 年新增）", None)):
        PP.cell(r, 1, lab).font = BOLD
        for y in YEARS:
            c = YC[y]
            if key == "camp":
                f = f'=SUMIFS({rng("L")},{rng("Q")},1,{rng("P")},{y})'
            elif key == "addchk":
                f = f"={c}{AGG['res']}+{c}{AGG['camp']}"
            elif y == 2023 and key in ("res", "robs", "runi"):
                f = 0
            else:
                f = f_(c)
            PP.cell(r, YEARS.index(y) + 4, f).number_format = "0.000"
        AGG[key] = r; r += 1
    r += 1
    PP.cell(r, 1, "C. 均勻型事件的年底上線比例矩陣（列＝年底 z，欄＝事件年 y；比例＝事件年內 asof 前部分 ＋ asof 後部分延後 d 年後在 z 年底前的部分）").font = BOLD; r += 1
    FRAC = {}
    for D in DLY:
        hdr(PP, r, [f"延遲情景 {D}", "年底 z", ""] + [""] + [str(y) for y in OUTY]); r += 1
        FRAC[D] = {}
        for z in YEARS:
            PP.cell(r, 1, f"{D}：{z} 年底"); PP.cell(r, 2, z)
            for y in OUTY:
                p1 = f"MAX(0,MIN({z}+1,{y}+MAX(0,MIN(1,asof-{y})))-{y})"
                def p2(d):
                    return f"MAX(0,MIN({z}+1,{y}+1+{d})-(MAX({y},asof)+{d}))"
                if D == "D0":
                    f = f"={p1}+{p2('0')}"
                elif D == "D1":
                    f = f"={p1}+{p2('dly_d1_m/12')}"
                else:
                    f = f"={p1}+(1-dly_d2_bind)*{p2('dly_d1_m/12')}+dly_d2_bind*{p2('(dly_d1_m+dly_d2_m)/12')}"
                PP.cell(r, YEARS.index(y) + 4, f).number_format = "0.00"
            FRAC[D][z] = r; r += 1
        r += 1
    PP.cell(r, 1, "D. 年底上線 TK GW（延遲情景）").font = BOLD; r += 1
    hdr(PP, r, ["項目", "情景", ""] + [str(y) for y in YEARS]); r += 1
    PA = {}
    def prow(key, lab, D, fz):
        nonlocal r
        PP.cell(r, 1, lab); PP.cell(r, 2, D)
        for z in YEARS:
            cc = PP.cell(r, YEARS.index(z) + 4, fz(z)); cc.number_format = "0.000"
        PA[key] = r; r += 1
    prow("stkc", "新納入存量累計", "", lambda z: f"=SUM($D${AGG['stk']}:{YC[z]}{AGG['stk']})")
    prow("obsc", "觀測型殘差累計（不延遲）", "", lambda z: f"=SUM($D${AGG['robs']}:{YC[z]}{AGG['robs']})")
    for D in DLY:
        prow("uni" + D, "均勻型殘差（延遲後）", D, lambda z, D=D: f"=SUMPRODUCT($E${AGG['runi']}:$M${AGG['runi']},$E${FRAC[D][z]}:$M${FRAC[D][z]})")
    prow("campD0", "園區（延遲後）", "D0", lambda z: f"=SUMPRODUCT({rng('L')},{rng('Q')},({rng('M')}<={z}+1)*1)")
    prow("campD1", "園區（延遲後）", "D1", lambda z: f"=SUMPRODUCT({rng('L')},{rng('Q')},({rng('N')}<={z}+1)*1)")
    prow("campD2", "園區（延遲後）", "D2", lambda z: f"=SUMPRODUCT({rng('L')},{rng('Q')},(1-dly_d2_bind)*({rng('N')}<={z}+1)+dly_d2_bind*({rng('O')}<={z}+1))")
    for D in DLY:
        prow("ye" + D, "年底上線 TK GW", D, lambda z, D=D: f"={YC[z]}{PA['stkc']}+{YC[z]}{PA['obsc']}+{YC[z]}{PA['uni' + D]}+{YC[z]}{PA['camp' + D]}")
        PP.cell(PA["ye" + D], 1).font = BOLD
        for z in YEARS: PP.cell(PA["ye" + D], YEARS.index(z) + 4).fill = KEY
    for D in DLY:
        prow("add" + D, "年新增 TK GW（淨，不含新納入存量）", D,
             lambda z, D=D: "=0" if z == 2023 else f"={YC[z]}{PA['ye' + D]}-{YC[z - 1]}{PA['ye' + D]}-{YC[z]}{AGG['stk']}")
    for D in ("D1", "D2"):
        prow("sno" + D, "已出貨未上線 TK GW（相對 D0）", D, lambda z, D=D: f"={YC[z]}{PA['yeD0']}-{YC[z]}{PA['ye' + D]}")
    PP_LAST = r
    PP.column_dimensions["A"].width = 40; PP.column_dimensions["B"].width = 26
    for col in "CDEFGHIJKLMNOPQ": PP.column_dimensions[col].width = 9

    # ================= Deploy：需求情景 × 延遲情景 =================
    DP["A1"] = ("Deploy — 市場總 GW（年底、年平均、年新增）× 需求情景 A／B／C × 延遲 D0／D1／D2。需求情景只改年新增：2024–2027 各情景＝Pipeline D0；"
                "2028–29＝Pipeline D0 年新增 × SIF_UnitsCoef_*；2030–32＝前年 ×（1＋g），g 由 2029 年增率線性收斂到 g_lt（2032 達 g_lt）。"
                "情景差額（相對 Pipeline D0 年新增）視為年內均勻上線，套用 Pipeline C 節的延遲比例")
    hdr(DP, 3, ["項目", "情景", "延遲"] + [str(y) for y in YEARS] + ["說明"])
    r = 4; DR = {}
    def drow(key, lab, s, D, fy, fmt="0.000", fill=None):
        nonlocal r
        DP.cell(r, 1, lab); DP.cell(r, 2, s); DP.cell(r, 3, D)
        for y in YEARS:
            cc = DP.cell(r, YEARS.index(y) + 4, fy(y)); cc.number_format = fmt
            if fill: cc.fill = fill
        DR[key] = r; r += 1
    drow("pipeadd", "Pipeline D0 年新增（淨）", "", "D0", lambda y: f"=Pipeline!{YC[y]}{PA['addD0']}")
    drow("stk", "新納入存量", "", "", lambda y: f"=Pipeline!{YC[y]}{AGG['stk']}")
    r += 1
    for s in SCEN:
        DP.cell(r, 1, f"需求情景 {s}").font = BOLD; r += 1
        gk = "g" + s; ak = "add" + s
        DR[gk] = r + 1  # 先佔位，add 列在前
        def fadd(y, s=s):
            c = YC[y]
            if y == 2023: return "=0"
            if y <= 2027: return f"={c}{DR['pipeadd']}"
            if y <= 2029: return f"={c}{DR['pipeadd']}*{sv('SIF_UnitsCoef_' + s, y)}"
            return f"={YC[y - 1]}{r}*(1+{c}{r + 1})"
        drow(ak, "年新增（需求情景，D0）", s, "D0", fadd)
        def fg(y, s=s):
            if y < 2029: return '=""'
            if y == 2029: return f"=IFERROR({YC[2029]}{DR[ak]}/{YC[2028]}{DR[ak]}-1,g_lt)"
            return f"=${YC[2029]}${r}+(g_lt-${YC[2029]}${r})*({y}-2029)/3"
        drow(gk, "年新增成長率 g（2029 起；2030–32 線性收斂到 g_lt）", s, "", fg, "0.0%")
        drow("dl" + s, "情景差額（相對 Pipeline D0 年新增）", s, "", lambda y, s=s: "=0" if y == 2023 else f"={YC[y]}{DR['add' + s]}-{YC[y]}{DR['pipeadd']}")
        for D in DLY:
            def fye(z, s=s, D=D):
                return f"=Pipeline!{YC[z]}{PA['ye' + D]}+SUMPRODUCT($E${DR['dl' + s]}:$M${DR['dl' + s]},Pipeline!$E${FRAC[D][z]}:$M${FRAC[D][z]})"
            drow(f"ye{s}{D}", "年底上線 TK GW", s, D, fye, fill=KEY)
            DP.cell(DR[f"ye{s}{D}"], 1).font = BOLD
        for D in DLY:
            drow(f"avg{s}{D}", "年平均 TK GW", s, D, lambda y, s=s, D=D: '=""' if y == 2023 else f"=({YC[y - 1]}{DR['ye' + s + D]}+{YC[y]}{DR['ye' + s + D]})/2")
        for D in DLY:
            drow(f"na{s}{D}", "年新增 TK GW（淨，不含新納入存量）", s, D,
                 lambda y, s=s, D=D: "=0" if y == 2023 else f"={YC[y]}{DR['ye' + s + D]}-{YC[y - 1]}{DR['ye' + s + D]}-{YC[y]}{DR['stk']}")
        for D in ("D1", "D2"):
            drow(f"sno{s}{D}", "已出貨未上線 TK GW（相對 D0）", s, D, lambda y, s=s, D=D: f"={YC[y]}{DR['ye' + s + 'D0']}-{YC[y]}{DR['ye' + s + D]}")
        r += 1
    DP.cell(r, 1, "對外發布：A-D0、B-D1、C-D2 三條與基準 A-D1（規劃書第 4 節）；Interface 全部 9 組皆輸出").font = BOLD
    DP_LAST = r + 1
    DP.column_dimensions["A"].width = 46
    for col in "BC": DP.column_dimensions[col].width = 6
    DP.column_dimensions["N"].width = 30

    # ================= C0：由下而上 vs 第三方 =================
    C0["A1"] = ("C0 — 市場總 GW：由下而上（Pipeline D0＝Σ 營運者）vs 第三方總量（每來源一列；口徑換算後、比值、容許帶、狀態）。"
                "第三方只作佐證、不校準；雙方皆 Verified 才判『對得上／偏差』，否則『證據不足（帶內／帶外）』。重複計算以營運者口徑加總一次處理（Matrix）")
    hdr(C0, 3, ["項目", "來源鍵", "口徑"] + [str(y) for y in YEARS] + ["說明"])
    r = 4
    C0.cell(r, 1, "由下而上 年底 TK GW（Pipeline D0）").font = BOLD; C0.cell(r, 2, "BU")
    for y in YEARS: C0.cell(r, YEARS.index(y) + 4, f"=Pipeline!{YC[y]}{PA['yeD0']}").number_format = "0.000"
    BU = r; r += 1
    C0.cell(r, 1, "有值玩家數（共 %d 列）" % len(PK))
    for y in YEARS: C0.cell(r, YEARS.index(y) + 4, f"=Players!{YC[y]}{TOT['cnt']}")
    r += 2
    C0R = []
    for o in seed["thirdparty"]:
        ir = INROW[o["key"]]
        C0.cell(r, 1, o["label"]).font = BOLD; C0.cell(r, 2, o["key"]); C0.cell(r, 3, f"=In!C{ir}")
        r += 1
        C0.cell(r, 1, "　原值（GW）")
        for y in YEARS: C0.cell(r, YEARS.index(y) + 4, f'=IF(In!{IC[y]}{ir}="","",In!{IC[y]}{ir})').font = GREEN
        rv = r; r += 1
        C0.cell(r, 1, "　換算後 TK GW（IT／未說明 ÷ g）")
        for y in YEARS:
            c = YC[y]
            C0.cell(r, YEARS.index(y) + 4, f'=IF({c}{rv}="","",IF(In!$C${ir}="設施",{c}{rv},{c}{rv}/Players!{c}${GROW}))').number_format = "0.000"
        rc = r; r += 1
        C0.cell(r, 1, "　比值（由下而上 ÷ 第三方）")
        for y in YEARS:
            c = YC[y]
            C0.cell(r, YEARS.index(y) + 4, f'=IF(OR({c}{rc}="",{c}{rc}=0),"",{c}${BU}/{c}{rc})').number_format = "0.00"
        rr_ = r; r += 1
        C0.cell(r, 1, "　狀態")
        for y in YEARS:
            c = YC[y]
            C0.cell(r, YEARS.index(y) + 4,
                    f'=IF({c}{rr_}="","無資料",IF(AND(In!$E${ir}="Verified",FALSE),IF(ABS({c}{rr_}-1)<=tolband,"對得上","偏差"),"證據不足（"&IF(ABS({c}{rr_}-1)<=tolband,"帶內","帶外")&"）"))')
        C0.cell(r, 14, "由下而上為 Derived（混合標記），v0.1 一律不判『對得上』")
        C0R.append(r); r += 2
    C0_LAST = r
    C0.column_dimensions["A"].width = 44; C0.column_dimensions["N"].width = 40
    for y in YEARS: C0.column_dimensions[YC[y]].width = 13

    # ================= C1：資本支出對帳 =================
    C1["A1"] = ("C1 — 資本支出對帳：GW 帳推算＝Σ 世代 新增 GW × TK IF_CapexIT（上線年入帳）＋ IF_CapexFacility（提前 lag_bldg 月入帳）；"
                "非 NVIDIA 部分＝持有比 IF_NonNV_Hold_* × VR200 每 GW；＋非 AI 資本支出（In Y3）；曆年按財年結束月換算成財年 vs 個股模型（In Y2）。"
                "第一個已知年的存量不計資本支出；2031–32 廠房缺次年上線量，偏低")
    hdr(C1, 3, ["項目", "玩家", "單位"] + [str(y) for y in YEARS] + ["可能原因／說明"])
    r = 4
    C1.cell(r, 1, "廠房提前入帳權重（lag_bldg 月 → l＝月/12；0≤l≤2）").font = BOLD; r += 1
    C1.cell(r, 1, "w0（上線當年）"); C1.cell(r, 4, "=MAX(0,1-lag_bldg/12)").number_format = "0.00"; W0 = f"$D${r}"; r += 1
    C1.cell(r, 1, "w1（前一年）"); C1.cell(r, 4, "=MAX(0,MIN(lag_bldg/12,2-lag_bldg/12))").number_format = "0.00"; W1 = f"$D${r}"; r += 1
    C1.cell(r, 1, "w2（前兩年）"); C1.cell(r, 4, "=MAX(0,MIN(1,lag_bldg/12-1))").number_format = "0.00"; W2 = f"$D${r}"; r += 1
    C1.cell(r, 1, "NVIDIA 每 TK GW IT capex（世代加權，US$B）")
    for y in YEARS:
        c = YC[y]
        C1.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexIT_{g}" for g in GENS)).number_format = "0.00"
    NVIT = r; r += 1
    C1.cell(r, 1, "NVIDIA 每 TK GW 廠房 capex（世代加權，US$B）")
    for y in YEARS:
        c = YC[y]
        C1.cell(r, YEARS.index(y) + 4, "=" + "+".join(f"Players!{c}{MIXR[g]}*TK_CapexFacility_{g}" for g in GENS)).number_format = "0.00"
    NVFAC = r; r += 2
    PLAT = {p["key"]: p["nnv_plat"] for p in P}
    REASON = {"msft": "未含融資租賃（10-K 另計）；租用 neocloud 容量（約 1.1 GW）不在資本支出；非 AI 待證據；財年 MW 置於曆年（GL01-04）",
              "googl": "TPU 每 GW 成本以持有比近似；含非 AI（搜尋、YouTube 等）；Epoch 推估 MW",
              "amzn": "資本支出含零售物流（非 AWS）；Trainium 每 GW 成本以持有比近似",
              "orcl": "Stargate 站點多為承租（房東資本支出）；財年 5 月；MW 為模型路徑",
              "crwv": "建物多為房東（Core Scientific、Galaxy）承租；公司資本支出含預付",
              "nbis": "自建＋承租混合；connected 口徑未說明",
              "wyfi": "共識資本支出未取得"}
    C1R = {}
    for p in PK:
        rows = PR[p]; plat = PLAT[p]
        cr = INROW[f"cap_{p}"]; nr = INROW[f"nai_{p}"]
        fye = FYE.get(p, "12")
        C1.cell(r, 1, PN[p]).font = BOLD; C1.cell(r, 2, p); C1.cell(r, 14, REASON.get(p, "個股模型待證據"))
        r += 1
        ra = r
        C1.cell(r, 1, "　IT 資本支出（上線年）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4, f"=Players!{c}{rows['add']}*((1-Players!{c}{rows['nnv']})*{c}${NVIT}+Players!{c}{rows['nnv']}*TK_NonNV_Hold_{plat}*TK_CapexIT_VR200)").number_format = "0.0"
        r += 1; rb = r
        C1.cell(r, 1, "　廠房資本支出（按上線年）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4, f"=Players!{c}{rows['add']}*((1-Players!{c}{rows['nnv']})*{c}${NVFAC}+Players!{c}{rows['nnv']}*TK_NonNV_Hold_{plat}*TK_CapexFacility_VR200)").number_format = "0.0"
        r += 1; rc = r
        C1.cell(r, 1, "　廠房資本支出（入帳曆年，提前 lag_bldg）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            t = [f"{YC[y]}{rb}*{W0}"]
            if y + 1 in YC: t.append(f"{YC[y + 1]}{rb}*{W1}")
            if y + 2 in YC: t.append(f"{YC[y + 2]}{rb}*{W2}")
            C1.cell(r, YEARS.index(y) + 4, "=" + "+".join(t)).number_format = "0.0"
        r += 1; rd = r
        C1.cell(r, 1, "　GW 帳推算 AI 資本支出（曆年）"); C1.cell(r, 3, "US$B")
        for y in YEARS: C1.cell(r, YEARS.index(y) + 4, f"={YC[y]}{ra}+{YC[y]}{rc}").number_format = "0.0"
        r += 1; re_ = r
        C1.cell(r, 1, "　GW 帳推算資本支出（財年；含非 AI）"); C1.cell(r, 3, "US$B")
        for y in YEARS:
            c = YC[y]
            prev = f"+{YC[y - 1]}{rd}*(12-{fye})/12" if y - 1 in YC else ""
            C1.cell(r, YEARS.index(y) + 4, f"={c}{rd}*{fye}/12{prev}+N(In!{IC[y]}{nr})").number_format = "0.0"
        r += 1; rf = r
        C1.cell(r, 1, "　個股模型資本支出（財年；In Y2）"); C1.cell(r, 3, "US$B")
        for y in YEARS: C1.cell(r, YEARS.index(y) + 4, f'=IF(In!{IC[y]}{cr}="","",In!{IC[y]}{cr})').font = GREEN
        r += 1; rg = r
        C1.cell(r, 1, "　比值（GW 帳 ÷ 個股模型）"); C1.cell(r, 3, "x")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4, f'=IF(OR({c}{rf}="",N({c}{rf})=0),"",{c}{re_}/{c}{rf})').number_format = "0.00"
        r += 1
        C1.cell(r, 1, "　狀態"); C1.cell(r, 3, "—")
        for y in YEARS:
            c = YC[y]
            C1.cell(r, YEARS.index(y) + 4,
                    f'=IF({c}{rg}="","無個股值",IF({c}{re_}=0,"GW 帳無新增",IF(AND(In!$E${cr}="Verified",FALSE),IF(ABS({c}{rg}-1)<=tolband,"對得上","偏差"),"證據不足（"&IF(ABS({c}{rg}-1)<=tolband,"帶內","帶外")&"）")))')
        C1R[p] = dict(gl=re_, model=rf, ratio=rg, st=r); r += 2
    C1_LAST = r
    C1.column_dimensions["A"].width = 44; C1.column_dimensions["N"].width = 70
    for y in YEARS: C1.column_dimensions[YC[y]].width = 13

    # ================= Compare：GW 帳 vs Siliconomics 成本等值 GW =================
    CP["A1"] = ("Compare — GW 帳年新增 TK GW（實體；營運者加總）vs Siliconomics SIF_DeployGW_A/B/C（成本等值 GW：AI IT 資本支出 ÷ 每 GW IT capex）。"
                "兩者口徑不同，只並列、不校準")
    hdr(CP, 3, ["項目", "情景", "延遲"] + [str(y) for y in YEARS] + ["說明"])
    r = 4; CPR = {}
    for s in SCEN:
        CP.cell(r, 1, f"Siliconomics 部署 GW（成本等值）"); CP.cell(r, 2, s)
        for y in (2026, 2027, 2028, 2029): CP.cell(r, YEARS.index(y) + 4, f"={sv('SIF_DeployGW_' + s, y)}").number_format = "0.00"
        CP.cell(r, 14, "SIL_In（Siliconomics 標『IT 電力』，除數為每 TK GW IT capex，實為 TK GW；規劃書第 6 節第 2 點）")
        CPR["sil" + s] = r; r += 1
        for D in ("D0", "D1"):
            CP.cell(r, 1, "GW 帳 年新增 TK GW（淨）"); CP.cell(r, 2, s); CP.cell(r, 3, D)
            for y in YEARS: CP.cell(r, YEARS.index(y) + 4, f"=Deploy!{YC[y]}{DR['na' + s + D]}").number_format = "0.00"
            CPR[f"gl{s}{D}"] = r; r += 1
        CP.cell(r, 1, "差額（GW 帳 D0 − Siliconomics）"); CP.cell(r, 2, s)
        for y in (2026, 2027, 2028, 2029):
            c = YC[y]; CP.cell(r, YEARS.index(y) + 4, f"={c}{CPR['gl' + s + 'D0']}-{c}{CPR['sil' + s]}").number_format = "0.00"
        CPR["d" + s] = r; r += 1
        CP.cell(r, 1, "比值（GW 帳 D0 ÷ Siliconomics）"); CP.cell(r, 2, s)
        for y in (2026, 2027, 2028, 2029):
            c = YC[y]; CP.cell(r, YEARS.index(y) + 4, f'=IF(N({c}{CPR["sil" + s]})=0,"",{c}{CPR["gl" + s + "D0"]}/{c}{CPR["sil" + s]})').number_format = "0.00"
        r += 2
    for t in ("差額原因（v0.1）：",
              "1. 覆蓋：GW 帳 v0.1 只有 repo 內已有 MW 的 7 家（MSFT、GOOGL、AMZN、ORCL、CRWV、NBIS、WYFI）；Meta、xAI、中國、主權、其他新雲待證據，GW 帳年新增大幅偏低。",
              "2. 公告路徑缺漏：多數玩家 2027 年後無公告值，Players 持平沿用 → 年新增為 0（CoreWeave 2030 目標內插除外）。",
              "3. 成本等值 vs 實體：Siliconomics 以 NVIDIA 每 GW 成本換算 ASIC 支出，ASIC 實體 GW 被低估約 0.3–4.3 GW（2026；規劃書第 6 節第 1 點）。",
              "4. 出貨年 ≠ 上線年：Siliconomics 為出貨（資本支出）年，GW 帳 D0 為公告上線年；D1／D2 再延後。"):
        CP.cell(r, 1, t); r += 1
    CP_LAST = r
    CP.column_dimensions["A"].width = 40; CP.column_dimensions["N"].width = 60

    # ================= Interface：GIF_ =================
    IF["A1"] = ("Interface — 下游唯一連結點（GIF_ 具名範圍；經 gwledger/tools/export_snapshot.py 匯成快照）。每格一個名稱：<前綴>_<yyyy>，2024–2032。"
                "單位一律 TK GW（設施電力）")
    IF["A2"] = f"版本：GW 帳 {a.version}（{a.date}）；Tokenomics {src['version']}；Siliconomics {ssrc['version']}"
    hdr(IF, 4, ["名稱前綴", "項目", "單位"] + [str(y) for y in OUTY] + ["來源位置"])
    r = 5; GIF = []
    def irow(prefix, lab, sheet, row):
        nonlocal r
        IF.cell(r, 1, prefix); IF.cell(r, 2, lab); IF.cell(r, 3, "GW")
        for j, y in enumerate(OUTY):
            c = IF.cell(r, 4 + j, f"={sheet}!{YC[y]}{row}"); c.font = GREEN; c.number_format = "0.000"
            nm_ = f"{prefix}_{y}"
            name(wb, nm_, f"Interface!${get_column_letter(4 + j)}${r}"); GIF.append(nm_)
        IF.cell(r, 13, f"{sheet}!{YC[2024]}{row}:{YC[2032]}{row}")
        r += 1
    IF.cell(r, 1, "A. 市場總 GW（需求情景 × 延遲）").font = BOLD; r += 1
    for s in SCEN:
        for D in DLY:
            irow(f"GIF_MktGW_YE_{s}_{D}", f"年底上線（{s}／{D}）", "Deploy", DR[f"ye{s}{D}"])
    for s in SCEN:
        for D in DLY:
            irow(f"GIF_MktGW_Avg_{s}_{D}", f"年平均（{s}／{D}）", "Deploy", DR[f"avg{s}{D}"])
    for s in SCEN:
        for D in DLY:
            irow(f"GIF_MktGW_Add_{s}_{D}", f"年新增（淨；{s}／{D}）", "Deploy", DR[f"na{s}{D}"])
    IF.cell(r, 1, "B. 已出貨未上線（相對 D0）").font = BOLD; r += 1
    for s in SCEN:
        for D in ("D1", "D2"):
            irow(f"GIF_ShippedNotOnline_{s}_{D}", f"已出貨未上線（{s}／{D}）", "Deploy", DR[f"sno{s}{D}"])
    IF.cell(r, 1, "C. 玩家年底上線 TK GW（甲：公告路徑；營運者口徑；空白＝未知）").font = BOLD; r += 1
    for p in PK:
        irow(f"GIF_PlayerGW_YE_{p}", PN[p], "Players", PR[p]["tk"])
    IF_LAST = r
    IF.column_dimensions["A"].width = 30; IF.column_dimensions["B"].width = 44; IF.column_dimensions["M"].width = 22

    # ================= Checks =================
    CK["A1"] = "Checks — 恆等式與治理（ERROR 計入 GGOV_Errors，須為 0；WARN 計入 GGOV_Warn，列為待補）"
    hdr(CK, 3, ["ID", "項目", "本庫值", "參照值", "差", "判讀", "性質", "說明"])
    r = 4
    maxr = {"In": IN_LAST + 2, "TK_In": TK_LAST + 2, "SIL_In": 60, "Players": PL_LAST + 2, "Matrix": MX_LAST + 2, "Pipeline": PP_LAST + 2,
            "Deploy": DP_LAST + 2, "C0": C0_LAST + 2, "C1": C1_LAST + 2, "Compare": CP_LAST + 2, "Interface": IF_LAST + 2}
    for sh, mr in maxr.items():
        CK.cell(r, 1, f"E-{sh}"); CK.cell(r, 2, f"{sh} 頁錯誤格數")
        CK.cell(r, 3, f"=SUMPRODUCT(ISERROR({sh}!A1:T{mr})*1)")
        CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "結構"); r += 1
    for y in (2025, 2026, 2027):
        m = MXR[y]
        CK.cell(r, 1, f"K1-{y}a"); CK.cell(r, 2, f"Matrix {y}：Σ|列和 − Players|")
        CK.cell(r, 3, f"=SUMPRODUCT(ABS(Matrix!J{m['r0']}:J{m['r1']}-Matrix!K{m['r0']}:K{m['r1']}))")
        CK.cell(r, 6, f'=IF(C{r}<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
        CK.cell(r, 1, f"K1-{y}b"); CK.cell(r, 2, f"Matrix {y}：欄和合計（含未分配）＝總量＝Players 市場合計")
        CK.cell(r, 3, f"=SUM(Matrix!C{m['sum']}:I{m['sum']})"); CK.cell(r, 4, f"=Players!{YC[y]}{TOT['ye']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for y in YEARS:
        CK.cell(r, 1, f"K2-{y}"); CK.cell(r, 2, f"Pipeline D0 {y} 年底＝Players 市場合計（容差 0.01 GW）")
        CK.cell(r, 3, f"=Pipeline!{YC[y]}{PA['yeD0']}"); CK.cell(r, 4, f"=Players!{YC[y]}{TOT['ye']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<=0.01,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for y in OUTY:
        CK.cell(r, 1, f"K3-{y}"); CK.cell(r, 2, f"Pipeline {y} 年底 D0 ≥ D1 ≥ D2")
        CK.cell(r, 3, f"=Pipeline!{YC[y]}{PA['yeD1']}"); CK.cell(r, 4, f"=Pipeline!{YC[y]}{PA['yeD0']}")
        CK.cell(r, 6, f'=IF(AND(Pipeline!{YC[y]}{PA["yeD0"]}>=Pipeline!{YC[y]}{PA["yeD1"]}-1E-9,Pipeline!{YC[y]}{PA["yeD1"]}>=Pipeline!{YC[y]}{PA["yeD2"]}-1E-9),"OK","ERROR")')
        CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K3-2032eq"); CK.cell(r, 2, "延遲只移動時點：D0 與 D2 的 2032 年底差 ≤ 2033 後才上線的部分（2032 均勻型殘差 ×（1 − 2032 年底 D2 比例）＋園區）")
    CK.cell(r, 3, f"=Pipeline!M{PA['yeD0']}-Pipeline!M{PA['yeD2']}")
    CK.cell(r, 4, f"=SUMPRODUCT(ABS(Pipeline!E{AGG['runi']}:M{AGG['runi']}),1-Pipeline!E{FRAC['D2'][2032]}:M{FRAC['D2'][2032]})+SUMPRODUCT(Pipeline!L{CR0}:L{CR1},Pipeline!Q{CR0}:Q{CR1})")
    CK.cell(r, 6, f'=IF(C{r}<=D{r}+1E-9,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    for y in YEARS:
        CK.cell(r, 1, f"K4-{y}"); CK.cell(r, 2, f"{y} 世代占比合計＝1")
        CK.cell(r, 3, f"=Players!{YC[y]}{MIXSUM}"); CK.cell(r, 4, 1)
        CK.cell(r, 6, f'=IF(ABS(C{r}-1)<1E-9,"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    for y in OUTY:
        CK.cell(r, 1, f"K5-{y}"); CK.cell(r, 2, f"{y} 殘差＋園區＝Players 年新增")
        CK.cell(r, 3, f"=Pipeline!{YC[y]}{AGG['addchk']}"); CK.cell(r, 4, f"=Players!{YC[y]}{TOT['add']}")
        CK.cell(r, 5, f"=C{r}-D{r}"); CK.cell(r, 6, f'=IF(ABS(E{r})<1E-9,"OK","ERROR")'); CK.cell(r, 7, "恆等式"); r += 1
    for k, lo, hi in (("lag_bldg", 0, 24), ("dly_d1_m", 0, 36), ("dly_d2_m", 0, 36), ("dly_d2_bind", 0, 1), ("tolband", 0, 1)):
        for suf in ("", "_lo", "_hi"):
            nm_ = k + suf
            CK.cell(r, 1, f"K6-{nm_}"); CK.cell(r, 2, f"{nm_} 在 {lo}–{hi}")
            CK.cell(r, 3, f"={nm_}"); CK.cell(r, 4, f"{lo}–{hi}")
            CK.cell(r, 6, f'=IF(AND(C{r}>={lo},C{r}<={hi}),"OK","ERROR")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K7"); CK.cell(r, 2, "未分園區殘差 < 0 的格數（園區合計超過玩家路徑）")
    CK.cell(r, 3, f"=COUNTIF(Pipeline!E{RES[PK[0]]}:M{RES[PK[-1]]},\"<-1E-9\")")
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證"); r += 1
    CK.cell(r, 1, "K8"); CK.cell(r, 2, "Tokenomics 快照含 IF_FacilityGW（有則 g 應改讀 TK）")
    CK.cell(r, 3, 1 if HAS_FACGW else 0)
    CK.cell(r, 6, '=IF(C%d=1,"OK","WARN")' % r); CK.cell(r, 7, "治理（TK）")
    CK.cell(r, 8, "v21.7 快照無 IF_FacilityGW；g 暫用規劃書 919.8 MW IT／TK GW（Requests R1）"); r += 1
    CK.cell(r, 1, "K9"); CK.cell(r, 2, "Deploy：各需求情景 2032 年底 D0 ≥ D1 ≥ D2（情景差額為負時可能不成立）")
    CK.cell(r, 3, "=" + "+".join(f"(Deploy!M{DR['ye' + s + 'D0']}<Deploy!M{DR['ye' + s + 'D1']}-1E-9)+(Deploy!M{DR['ye' + s + 'D1']}<Deploy!M{DR['ye' + s + 'D2']}-1E-9)" for s in SCEN))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "驗證"); r += 1
    a0, a1 = 5, IN_LAST
    CK.cell(r, 1, "G1"); CK.cell(r, 2, "In 頁輸入列缺標記")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}<>"")*(In!E{a0}:E{a1}=""))'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G2"); CK.cell(r, 2, "In 頁輸入列缺來源 ID")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}<>"")*(In!F{a0}:F{a1}=""))'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G3"); CK.cell(r, 2, "純量列標記含 Assumed／Analogy 但未給低／高")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}=1)*((ISNUMBER(SEARCH("Assumed",In!E{a0}:E{a1}))+ISNUMBER(SEARCH("Analogy",In!E{a0}:E{a1})))>0)*(((In!H{a0}:H{a1}="")+(In!J{a0}:J{a1}=""))>0))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G3b"); CK.cell(r, 2, "年度／分攤列標記含 Assumed 或 Analogy（年度列無低／高欄，v0.2 補區間列）")
    CK.cell(r, 3, f'=SUMPRODUCT((In!S{a0}:S{a1}>1)*((ISNUMBER(SEARCH("Assumed",In!E{a0}:E{a1}))+ISNUMBER(SEARCH("Analogy",In!E{a0}:E{a1})))>0))')
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理"); r += 1
    keys = list(INROW)
    CK.cell(r, 1, "G4"); CK.cell(r, 2, "In 頁鍵重複（建置時檢查）"); CK.cell(r, 3, len(keys) - len(set(keys)))
    CK.cell(r, 6, f'=IF(C{r}=0,"OK","ERROR")'); CK.cell(r, 7, "治理"); r += 1
    CK.cell(r, 1, "G5"); CK.cell(r, 2, "標『待證據』的輸入列數")
    CK.cell(r, 3, f'=COUNTIF(In!E{a0}:E{a1},"待證據")'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理")
    CK.cell(r, 8, "evidence_raw 複核後入 In，逐步歸零"); r += 1
    CK.cell(r, 1, "G6"); CK.cell(r, 2, "延遲參數仍為佔位值（說明含【佔位】）")
    CK.cell(r, 3, f'=COUNTIF(In!G{a0}:G{a1},"*【佔位】*")'); CK.cell(r, 6, f'=IF(C{r}=0,"OK","WARN")'); CK.cell(r, 7, "治理"); r += 1
    r += 1
    CK.cell(r, 1, "GGOV_Errors").font = BOLD; c = CK.cell(r, 3, f'=COUNTIF(F4:F{r - 2},"ERROR")'); c.fill = KEY
    name(wb, "GGOV_Errors", f"Checks!$C${r}"); r += 1
    CK.cell(r, 1, "GGOV_Warn").font = BOLD; c = CK.cell(r, 3, f'=COUNTIF(F4:F{r - 3},"WARN")'); c.fill = KEY
    name(wb, "GGOV_Warn", f"Checks!$C${r}")
    for col, w in zip("ABCDEFGH", (12, 70, 12, 12, 10, 8, 12, 60)):
        CK.column_dimensions[col].width = w

    # ================= Decisions、DB_Evidence、Sources、Requests =================
    DC["A1"] = "Decisions — 決策登錄（GL01-01 起）。判斷類由 Andy 決定；工程細節依預設執行者標『預設、可撤回』"
    hdr(DC, 3, ["決策 ID", "類別", "項目", "決議", "日期", "依據", "狀態"])
    decisions = list(seed["decisions"]) + [
        {"id": "GL01-11", "cat": "工程", "item": "Rubin Ultra 占比", "decision": "Units 顆數 ÷ TK 每 GW GPU 數換成 GW 占比（RU ÷（VR200＋RU））；2027 起 Blackwell 不計", "date": "2026-10-28", "quote": "工作單第 2 節", "status": "預設、可撤回"},
        {"id": "GL01-12", "cat": "待決", "item": "延遲參數 D1／D2、g_lt、lag_bldg、2023–25 NonNV 比例", "decision": "佔位值（In S1、S2；Y6 nnv_mkt）；待 B_delay.csv 與 Andy 判斷", "date": "2026-10-28", "quote": "chat 端建置指示", "status": "預設、可撤回"},
    ]
    r = 4
    for d in decisions:
        for j, k in enumerate(("id", "cat", "item", "decision", "date", "quote", "status"), 1):
            c = DC.cell(r, j, d.get(k)); c.alignment = WRAP
        if "預設" in str(d.get("status")):
            for j in range(1, 8): DC.cell(r, j).fill = PEND
        r += 1
    for col, w in zip("ABCDEFG", (9, 8, 24, 80, 12, 34, 14)):
        DC.column_dimensions[col].width = w
    EV["A1"] = "DB_Evidence — 證據登錄（gwledger/evidence_raw/*.csv 經 chat 端複核後整理入本頁；含不採納者）"
    hdr(EV, 3, ["ID", "日期", "主張（摘要）", "來源", "標記", "對應鍵", "本庫當時值", "新資訊值", "判定", "處理版本", "備註"])
    EV.cell(4, 1, "（v0.1 階段 1 骨架：尚無；evidence_raw 由查證代理產出，複核後交付再入本頁與 In）")
    SR["A1"] = "Sources — 來源 ID（In 頁 F 欄）"
    hdr(SR, 3, ["來源 ID", "路徑", "說明"])
    r = 4
    for s in seed["sources"]:
        SR.cell(r, 1, s["id"]); SR.cell(r, 2, s["path"]); SR.cell(r, 3, s["desc"]); r += 1
    for col, w in zip("ABC", (14, 70, 80)): SR.column_dimensions[col].width = w
    RQ["A1"] = "Requests — 需要 TK 或 Siliconomics 權威改值或新增名稱的請求（本庫不改其檔案；工作單第 4 節停止條件）"
    hdr(RQ, 3, ["ID", "對象", "請求", "原因", "狀態"])
    reqs = [("R1", "Tokenomics", "快照加入 IF_FacilityGW（每 TK GW 的 IT／設施比）或確認每 GW 口徑", "v21.7 快照註記『每 GW＝IT 關鍵電力』，與規劃書『TK GW＝設施電力；VR200 含 919.8 MW IT』不一致；g 暫用 0.9198", "待 Andy／TK"),
            ("R2", "Siliconomics", "Deploy 頁『部署 GW（IT 電力）』標籤改為 TK GW", "除數 IF_CapexIT 為每 TK GW；規劃書第 6 節第 2 點", "待 AI半導體"),
            ("R3", "Tokenomics", "NonNV 每 GW 顆數與每顆功率的 IF_ 名稱（v0.2 C2）", "規劃書第 5 節", "v0.2")]
    r = 4
    for q in reqs:
        for j, v in enumerate(q, 1): RQ.cell(r, j, v).alignment = WRAP
        r += 1
    for col, w in zip("ABCDE", (6, 14, 60, 70, 14)): RQ.column_dimensions[col].width = w

    # ================= README =================
    lines = [
        (f"GW 帳 {a.version} — 全市場 AI 資料中心 GW 唯一基準事實（{a.date}）", True),
        ("命題：市場總 GW 要與主要玩家 GW 對得上，再與資本支出、GPU／ASIC 顆數、記憶體、電廠與發電機組、電源散熱、AI 投資地圖個股逐環對得上；並加入與需求情景正交的基建延遲情景。", False),
        ("驅動 → 推導：In（玩家年底 MW、園區管線、延遲 D0–D2、世代組合、分攤）＋ TK_In（每 GW 係數）＋ SIL_In（顆數、情景係數）→ Players（玩家 × 年 TK GW）→ Matrix（營運者 × 使用者）→ Pipeline（上線事件 × 延遲）→ Deploy（需求 × 延遲 3×3）→ C0／C1／Compare（對帳）→ Interface（GIF_）→ Checks（GGOV_Errors＝0）。", False),
        ("單位：TK GW＝設施電力（VR200 每 TK GW 含 919.8 MW IT）；IT 讀法 ÷ g。口徑未說明者基準取 IT 讀法，Players 另列設施讀法。", False),
        ("頁面：README｜In 輸入｜TK_In、SIL_In 鏡像（綠字）｜Players 玩家 × 年｜Matrix 營運者 × 使用者｜Pipeline 園區管線 × 延遲｜Deploy 市場總 GW 需求 × 延遲｜C0 由下而上 vs 第三方｜C1 資本支出對帳｜Compare 與 Siliconomics 成本等值 GW｜Interface GIF_｜Checks｜Decisions｜DB_Evidence｜Sources｜Requests。", False),
        ("v0.1 範圍：C0、C1、延遲管線（D0–D2）、GIF_ 輸出；C2–C7 自 v0.2 起。階段 1（骨架）只放 repo 內已有、帶來源的值，其餘列留空並標『待證據』；延遲參數為佔位值。", False),
        ("顏色：藍字＝輸入；黑字＝公式；綠字＝跨頁或外部鏡像；淡黃底＝關鍵輸出；亮黃底＝待證據或預設待決。", False),
        ("Excel 優先：要改輸入，直接改 In 頁藍字；builder 重建時讀回（--base）。公式頁不要手改。下游只讀 Interface 的 GIF_ 名稱（經快照）。", False),
        ("版本命名：YYYYMMDD_GWLedger_vX.Y.xlsx；model/CURRENT 記錄現行檔名；data/CURRENT_SNAPSHOT 指向現行快照。", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = RD.cell(i, 1, t); c.alignment = WRAP
        if b: c.font = Font(bold=True, size=13)
    RD.column_dimensions["A"].width = 150

    wb.save(a.out)
    print(f"寫出 {a.out}：頁面 {len(wb.sheetnames)}、In 鍵 {len(INROW)}、GIF 名稱 {len(GIF)}、具名範圍 {len(wb.defined_names)}")


if __name__ == "__main__":
    main()
