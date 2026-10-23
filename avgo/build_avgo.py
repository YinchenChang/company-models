# Broadcom（AVGO）收支與估值模型 v0.1（AI 半導體 Project：運算層第二個樣板，ASIC＋網通）
# 用法：python3 avgo/build_avgo.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度 FY(t) 於 t 年 10 月底／11 月初結束（FY2026：2025-11-03～2026-11-01），10/12 落在日曆年 t，
# 所以 FY(t) 直接對照 CY(t)，不做平移（NVIDIA 需平移一年，AVGO 不需）。金額：十億美元（US$B）；EPS：美元；口徑：公司 non-GAAP。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "nvidia"))
import market as MEM
import cycle_position as CP
import accel_share as AS          # 加速器金額份額的可重用函式（NVIDIA 模型建立；AVGO 為 ASIC 側）

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
XP = ("tpu", "oth")                                   # Broadcom 的 XPU 來自這兩個市場層平台
PNAME = {k: nm for k, nm, *_ in MEM.UNITS}
PLAT = [k for k, *_ in MEM.UNITS]

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); SKH_F = only("skhynix/dist/*.xlsx"); NET_F = only("markets/network/dist/*.xlsx")
NV_F = only("nvidia/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)

ws_ = openpyxl.load_workbook(TSMC_F, data_only=True)["In"]
TS = {}
for r in ws_.iter_rows(min_row=4):
    k = r[0].value
    if k and (k.startswith("um2") or k.startswith("g29_") or k.startswith("u_")):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
for k, nm, u_, gb, tag, src in MEM.UNITS:     # 與 markets/memory 的 UNITS 必須一致（同一份）
    for i, y in enumerate((26, 27, 28)):
        assert abs(TS[f"u_{k}_{y}"][0] - u_[i]) < 1e-9, (k, y)
for s_ in SC:
    assert abs(TS[f"um28_{s_}"][0] - MEM.SCEN_UNITS[s_][0]) < 1e-9 and abs(TS[f"um29_{s_}"][0] - MEM.SCEN_UNITS[s_][1]) < 1e-9

wsn = openpyxl.load_workbook(NET_F, data_only=True)
NETB = None; REPB = None
for r in wsn["Mkt"].iter_rows(min_row=4):
    if r[0].value == "Broadcom 網通（本模型歸屬）":
        NETB = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
for r in wsn["In"].iter_rows(min_row=4):
    if r[0].value == "rep_avgo26":
        REPB = ((r[3].value, r[4].value, r[5].value), f"{rel(NET_F)}｜In!D{r[0].row}:F{r[0].row}")
assert NETB and REPB and all(v is not None for v in NETB[0])

wsk = openpyxl.load_workbook(SKH_F, data_only=True)
HBMP = {}
for r in wsk["Mkt"].iter_rows(min_row=4):
    lab = str(r[0].value or "")
    if lab.startswith("HBM 單價（"):
        HBMP[lab[-2]] = ([r[2 + j].value for j in range(5)], f"{rel(SKH_F)}｜Mkt!C{r[0].row}:G{r[0].row}")
assert sorted(HBMP) == list(SC), HBMP.keys()

wsv = openpyxl.load_workbook(NV_F, data_only=True)["In"]
NVP = {}
for r in wsv.iter_rows(min_row=4):
    if r[0].value in ("pm_bw", "pm_rb", "pm_ru", "pm_hp", "pa_amd", "pa_trn", "pa_tpu", "pa_oth"):
        NVP[r[0].value] = (r[4].value, f"{rel(NV_F)}｜In!E{r[0].row}")
assert len(NVP) == 8, NVP

PEAK, PEAK_SRC = CP.peak_years(ROOT)

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億美元（US$B）；每顆價格：千美元（US$K）；顆數：百萬顆；HBM：GB、US$/GB。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝Broadcom 8-K 新聞稿（SEC EDGAR、investors.broadcom.com）；公司對 FY2027、FY2028 的 AI 營收展望為 Interested-party，只放 Recon 對帳，不作校準"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
EDG = "https://www.sec.gov/Archives/edgar/data/1730168/"
Q3PR = "https://investors.broadcom.com/news-releases/news-release-details/broadcom-inc-announces-third-quarter-fiscal-year-2026-financial"
CALL = "法說轉述：Investing.com 逐字稿、BigGo Finance 摘要（2026-09-02；二手）"

sec("A. 加速器顆數（markets/memory 與 HBM、台積電、網通、NVIDIA 模型同一份；情境係數讀台積電模型成品，不另訂）")
for k, nm, u_, gb, tag, src in MEM.UNITS:
    for i, y in enumerate((26, 27, 28)):
        inp(f"u_{k}_{y}", f"{nm} 顆數 20{y}", "百萬顆", None, u_[i], None, tag, "markets/memory/market.py UNITS：" + src if i == 0 else "同上", f"台積電模型 {TS[f'u_{k}_{y}'][1]}（值相同，建置時檢查）" if i == 0 else "同上")
    inp(f"g29_{k}", f"{nm} 2029 顆數成長", "%", None, TS[f"g29_{k}"][0], None, "Assumed", TS[f"g29_{k}"][1], "markets/memory 無 2029；沿用台積電模型")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
for k in XP:
    gb = [g for kk, nm, u_, g, t, sr in MEM.UNITS if kk == k][0]
    for i, y in enumerate((26, 27, 28)):
        inp(f"gb_{k}_{y}", f"{PNAME[k]} 每顆 HBM 20{y}", "GB", None, gb[i], None, "Interested-party" if k == "tpu" else "Derived／Assumed", "markets/memory/market.py UNITS", "2029 沿用 2028" if i == 0 else "同上")

sec("B. ASIC 每顆規格（本專案為權威，規則第 13 條）：Broadcom 在各平台的顆數份額、每顆核心售價（不含 HBM）、HBM 由 Broadcom 代購的比例")
inp("sh_tpu_26", "Broadcom 占 Google TPU 顆數 2026", "%", 0.80, 0.90, 1.0, "Assumed", "Broadcom 為 TPU 主要設計夥伴（Ironwood／v7、v8i 量產；法說）；Google 與 MediaTek 另有合作（Counterpoint 2026-01，二手）", "各家份額未揭露：未找到")
inp("sh_tpu_27", "Broadcom 占 Google TPU 顆數 2027", "%", 0.60, 0.80, 0.90, "Assumed", "同上；MediaTek 版 TPU 放量的時點未揭露", "Anthropic 2027 年約 5 GW 的 TPU v8i（公司說法，Interested-party）屬 Broadcom")
inp("sh_tpu_28", "Broadcom 占 Google TPU 顆數 2028（2029 沿用）", "%", 0.50, 0.75, 0.90, "Assumed", "同上", "—")
inp("sh_oth_26", "Broadcom 占「其他 ASIC」顆數 2026", "%", 0.20, 0.35, 0.50, "Assumed", "其他 ASIC＝MTIA（Broadcom）、OpenAI 客製（Broadcom，2026 Q3 開始出貨）、Maia（非 Broadcom）、昇騰（華為）等", "XPU 客戶 6 家：Google、Anthropic、OpenAI、Meta＋2 家未具名（法說，二手）")
inp("sh_oth_27", "Broadcom 占「其他 ASIC」顆數 2027", "%", 0.40, 0.55, 0.70, "Assumed", "OpenAI 2027 年 1.3 GW、Meta 至 2028 年共 3 GW（公司說法，Interested-party）", "—")
inp("sh_oth_28", "Broadcom 占「其他 ASIC」顆數 2028（2029 沿用）", "%", 0.40, 0.55, 0.70, "Assumed", "OpenAI 2028 年 5 GW 以上（公司說法）", "—")
inp("core_tpu", "TPU 每顆核心售價（Broadcom 營收，不含 HBM；2026 基準）", "US$K", 6, 8, 11, "Assumed", "無：每顆 ASP 未揭露。NVIDIA 模型 pa_tpu（含 HBM）8–12–18K；192 GB × 2026 HBM 價約 2.8K", NVP["pa_tpu"][1] + "（含 HBM 的對照值）")
inp("core_oth", "其他 ASIC 每顆核心售價（Broadcom 營收，不含 HBM；2026 基準）", "US$K", 4, 6, 9, "Assumed", "無：未揭露。Maia 200 約為 NVIDIA 的 30–40%＝$9–12K（二手，含 HBM）作量級參照", NVP["pa_oth"][1] + "（含 HBM 的對照值）")
inp("f_hbm", "HBM 由 Broadcom 代購並計入營收的比例", "x", 0.3, 0.6, 1.0, "Assumed", "公司：毛利率下滑來自「記憶體含量高的 XPU」占比上升（Q3 FY26 法說，二手）→ 至少部分 HBM 經 Broadcom 轉售", "第二來源：未找到。客戶自購 HBM（consigned）的比例未揭露")
inp("m_hbm", "代購 HBM 的毛利率", "x", 0.0, 0.10, 0.20, "Assumed", "無：未揭露；轉售料件通常低加成", "第二來源：未找到")
inp("pfc", "2027 年起 XPU 核心售價係數（Sens 用；基準 1）", "x", 0.8, 1.0, 1.2, "Assumed", "無：新一代（TPU v8、MTIA 新世代）核心單價變化未揭露", "Sens 用 ±20%")

sec("C. 記憶體價格路徑（讀 skhynix 成品；唯讀；情境別）")
for s in SC:
    for j, y in enumerate((26, 27, 28, 29)):
        inp(f"hbm_{s}_{y}", f"HBM 混合單價 20{y}（{s}）", "US$/GB", None, HBMP[s][0][j + 1], None, "Derived", HBMP[s][1], "markets/memory：2026 BofA $14.6/GB；2027 +80% 三情境相同；2028–29 依情境" if (s == "A" and y == 26) else "同上")

sec("D. AI 網通（讀 markets/network 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"net_{y}", f"Broadcom 網通（本模型歸屬）20{y}", "US$B", None, NETB[0][j], None, "Derived", NETB[1], "廠商歸屬比例為 Assumed（markets/network In 頁 avgo_su、avgo_so、avgo_nic、avgo_opt）；2026 較實際推估低約 28%（見 Recon）" if y == 26 else "同上")
inp("rep_avgo26", "Broadcom AI 網通營收 FY2026（實際＋推估）", "US$B", REPB[0][0], REPB[0][1], REPB[0][2], "Derived", REPB[1], "Q3 FY26：XPU 占 AI 營收 73%（法說，二手）→ 網通約 4.5；FY2026 AI 57.6 × 約 30%")
inp("gm_net", "AI 網通毛利率", "x", 0.65, 0.72, 0.80, "Assumed", "無：公司不揭露產品別毛利率", "第二來源：未找到")

sec("E. 非 AI 半導體與基礎架構軟體（VMware 等）")
inp("gm_nai", "非 AI 半導體毛利率", "x", 0.60, 0.65, 0.70, "Assumed", "無：公司不揭露；半導體部門整體毛利率約 67%（Q3 FY26 法說 Q&A 更正值，二手）", "第二來源：未找到")
inp("gm_sw", "軟體毛利率（non-GAAP）", "x", 0.90, 0.93, 0.95, "Derived", "Q3 FY26：(non-GAAP 毛利 22.191 − 半導體 20.839 × 67%) ÷ 軟體 8.752 ≈ 94%（67% 為二手）", "低／高 Assumed")
inp("g_nai", "非 AI 半導體營收年增（FY2027 起）", "%", -0.05, 0.03, 0.08, "Assumed", "FY2026 Q3 約 4.1／季（半導體 20.8 − AI 16.7；公司）", "第二來源：未找到")
inp("g_sw27", "軟體營收年增 FY2027", "%", 0.08, 0.12, 0.18, "Assumed", "公司：Q3 FY26 軟體年增 29%、ARR 年增 15%；Q4 預估約 8.7（約 +25%）（法說，二手）", "第二來源：未找到")
inp("g_sw", "軟體營收年增 FY2028 起", "%", 0.04, 0.08, 0.12, "Assumed", "VMware 漲價效果遞減（Assumed）", "—")

sec("F. FY2026 錨點（Q1–Q3 實績＋Q4 指引）")
inp("q4_rev", "Q4 FY26 營收指引", "US$B", None, 34.8, None, "Interested-party", Q3PR + "：約 $34.8B（年增 93%）", "top1markets、BigGo（二手）同值")
inp("q4_ai", "Q4 FY26 AI 半導體營收", "US$B", None, 21.7, None, "Interested-party", "Q3 FY26 新聞稿 CEO 說明：$21.7B（年增 236%）", "BigGo、top1markets（二手）同值")
inp("q4_sw", "Q4 FY26 軟體營收", "US$B", 8.5, 8.7, 8.9, "Interested-party", CALL + "：約 $8.7B", "第二來源：未找到")
inp("q4_gm", "Q4 FY26 non-GAAP 毛利率", "%", 0.72, 0.73, 0.74, "Interested-party", CALL + "：約 73%（去年同期 78%）", "第二來源：未找到")
inp("q4_oim", "Q4 FY26 non-GAAP 營業利益率指引", "%", None, 0.66, None, "Interested-party", Q3PR + "：約 66%", "top1markets（二手）同值")
inp("q4_int", "Q4 FY26 non-GAAP 利息費用", "US$B", 0.65, 0.70, 0.75, "Derived", "Q3 FY26 0.703（公司）", "—")
inp("q4_oth", "Q4 FY26 non-GAAP 其他收益", "US$B", 0.05, 0.10, 0.15, "Derived", "Q1–Q3 0.118／0.118／0.098（公司）", "—")
inp("q4_tx", "Q4 FY26 non-GAAP 有效稅率", "%", 0.155, 0.16, 0.165, "Derived", "Q3 FY26 3.118 ÷ 19.490＝16.0%（公司）", "—")
inp("q4_sh", "Q4 FY26 non-GAAP 稀釋股數", "十億股", None, 4.94, None, "Derived", "Q3 FY26 4.937（公司）", "—")

sec("G. 營業費用、利息、稅、股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.05, 0.10, 0.30, "Derived", "FY2026 錨點營業費用年增 ÷ 營收年增（Q 頁，約 0.09）", "低／高 Assumed")
inp("g_opx_min", "營業費用年增下限", "%", 0.0, 0.03, 0.06, "Assumed", "無", "—")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 ±2 個百分點")
inp("int_f", "non-GAAP 利息費用（每年）", "US$B", 2.5, 2.8, 4.0, "Assumed", "FY2026 錨點約 2.9（Q 頁）；總債務 59.4（Q3 FY26，公司）", "高端：客戶融資若改由 Broadcom 舉債承擔（見 Recon F 節）")
inp("oth_f", "non-GAAP 其他收益（每年）", "US$B", 0.2, 0.4, 0.6, "Assumed", "FY2026 約 0.43（Q 頁）", "—")
inp("tx_f", "non-GAAP 有效稅率", "%", 0.15, 0.16, 0.18, "Derived", "FY2026 Q1–Q3 約 15.6–16.5%（公司）", "—")
inp("sh_g", "稀釋股數年變動", "%", -0.015, -0.005, 0.0, "Derived", "FY2025 4.943 → Q3 FY26 4.937（公司）", "低／高 Assumed")
inp("uk", "Broadcom XPU 顆數係數（Sens 用；FY2027 起）", "x", 0.9, 1.0, 1.1, "Assumed", "無", "±10%")

sec("H. 公司展望與共識（只作 Recon 對帳，不作校準）")
inp("co_ai26", "公司 FY2026 AI 營收展望", "US$B", None, 58.0, None, "Interested-party", CALL, "Q1–Q3 實績＋Q4 指引合計 57.6（Q 頁）")
inp("co_ai27", "公司 FY2027 AI 營收展望", "US$B", None, 115.0, None, "Interested-party", CALL + "：約 $115B；公司稱不會逐季更新", "JPMorgan：FY2027 超過 $150B（KuCoin 轉述，Interested-party）")
inp("co_ai28", "公司 FY2028 AI 營收展望", "US$B", None, 230.0, None, "Interested-party", CALL + "：約 $230B", "top1markets 2026-10-22（二手）同值")
inp("jpm_ai27", "JPMorgan FY2027 AI 營收估計", "US$B", None, 150.0, None, "Interested-party", "JPMorgan（KuCoin 2026-06-22 轉述）：超過 $150B", "—")
inp("cons_rev26", "共識 FY2026 營收", "US$B", None, 105.98, None, "Interested-party", "StockAnalysis（S&P Global，40 位分析師，2026-10-22）", "—")
inp("cons_eps26", "共識 FY2026 non-GAAP EPS", "US$", None, 11.66, None, "Interested-party", "同上（區間 11.35–12.38）", "—")
inp("cons_rev27", "共識 FY2027 營收", "US$B", None, 173.97, None, "Interested-party", "同上（「明年」欄，頁面未標年度）", "—")
inp("cons_eps27", "共識 FY2027 non-GAAP EPS", "US$", None, 19.40, None, "Interested-party", "同上", "—")
inp("jpm_u27", "2027 年 ASIC 占加速器出貨顆數（JPMorgan）", "%", None, 0.536, None, "Interested-party", "JPMorgan（KuCoin 轉述）：ASIC 12.5M ÷ 23.3M", "本專案市場層 ASIC 2027 12.5M 同一來源")
inp("cp_sh27", "2027 年 Broadcom 占 ASIC 設計夥伴份額（Counterpoint）", "%", None, 0.60, None, "Interested-party", "Counterpoint（Benzinga／Yahoo 2026-01-28 轉述；口徑未說明）", "JPMorgan：高階 ASIC 80–85%")

sec("I. 加速器每顆價格（金額份額計算用；讀 NVIDIA 模型成品，模組層級）")
for k, lab in (("pm_bw", "Blackwell 模組"), ("pm_rb", "Rubin 模組"), ("pm_ru", "Rubin Ultra 模組"), ("pm_hp", "Hopper 模組"), ("pa_amd", "AMD Instinct"), ("pa_trn", "AWS Trainium")):
    inp(k, f"{lab} 每顆", "US$K", None, NVP[k][0], None, "Derived", NVP[k][1], "NVIDIA 模型 In 頁（同一套價格；唯讀）")

sec("J. 價格、估值與客戶融資")
inp("px", "股價（2026-10-22 收盤）", "US$", None, 323.24, None, "Interested-party", "StockAnalysis（2026-10-22 收盤）", "YCharts 同值（2026-10-22 收盤 323.24）")
inp("pe_ttm_gaap", "近 12 月 GAAP P/E", "x", None, 41.26, None, "Interested-party", "StockAnalysis", "top1markets：近 12 月約 45x（10 月報導）")
inp("pe_fwd_sa", "前瞻 P/E（StockAnalysis）", "x", None, 18.61, None, "Interested-party", "StockAnalysis（口徑未說明）", "GuruFocus 17.39；valueinvesting.io 49.62（2026-10-22；口徑不明，可能為 GAAP 或舊預估，不採）")
inp("eps_q4_25", "Q4 FY25 non-GAAP EPS（近 12 月計算用）", "US$", None, 1.95, None, "Verified", EDG + "000173016825000116/avgo-11022025x8kxex99.htm", "—")
inp("pe_fair_andy", "Andy 先驗：合理 P/E", "x", None, 25, None, "Assumed", "先驗_Andy_AVGO_20261023 B9b（Drive；不在本 repo）", "只作對照")
inp("fin_anth", "Broadcom 對 Anthropic 放款上限（可轉股）", "US$B", None, 42, None, "Interested-party", "top1markets 2026-10-22 引 Reuters（Anthropic IPO 文件）", "Bloomberg：銀行團 $60B 中 $42B 優先擔保部分由 Broadcom 殘值支撐（二手）")
inp("fin_oai", "Broadcom 為 OpenAI 安排的融資", "US$B", None, 50, None, "Interested-party", "top1markets 引 WSJ：超過 $50B", "—")
inp("fin_lease", "一家客戶 AI 機架租賃擔保", "US$B", None, 29, None, "Interested-party", "top1markets 引 Broadcom 季報：最高約 $29B", "—")
for s in SC:
    inp(f"pk_{s}", f"情境 {s}：設備市場見頂年度（日曆年；循環位置規則用）", "年", None, PEAK[s], None, "Derived", PEAK_SRC[s], "共用規則（tools/cycle_position.py）：AVGO 的 FY(t) 對照日曆年 t（不平移）" if s == "A" else "同上")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績（十億美元；EPS 為美元；股數為十億股）。口徑：公司 non-GAAP（排除股票薪酬、無形資產攤銷、重組）。藍字＝輸入；黑字＝公式。實績取自 Broadcom 8-K 新聞稿（SEC EDGAR）"
Q["A2"] = "FY2026 錨點＝Q1–Q3 實績＋Q4 指引（F 欄；假設見 In 頁 F 節）。AI 營收：Q1 8.4、Q2 10.8、Q3 16.7（公司）、Q4 21.7（公司指引）；AI 網通取 markets/network 的實際推估 rep_avgo26，XPU＝AI − 網通"
hdr(Q, 3, ["項目", "FY2025A", "FY26 Q1（～2026-02-01）", "FY26 Q2（～05-03）", "FY26 Q3（～08-02）", "FY26 Q4（指引）", "FY2026 錨點", "標記", "來源1（公司）", "來源2／說明"])
S1 = ("FY2025、Q4 FY25：" + EDG + "000173016825000116/avgo-11022025x8kxex99.htm；Q1 FY26：" + EDG + "000173016826000011/avgo-02012026x8kxex99.htm；"
      "Q2 FY26：" + EDG + "000173016826000051/avgo-05032026x8kxex99.htm；Q3 FY26：" + Q3PR)
S2 = "交叉：Q3 FY26 稿的九個月合計（營收 71.089、non-GAAP 毛利 54.168、營業利益 47.849、淨利 38.631）＝Q1＋Q2＋Q3（下方檢查列）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
qrow(4, "營收", {"B": 63.887, "C": 19.311, "D": 22.187, "E": 29.591, "F": "=q4_rev"}, tag="Verified／Interested-party（Q4 指引）")
qrow(5, "  半導體解決方案", {"B": 36.858, "C": 12.515, "D": 15.009, "E": 20.839, "F": "=F4-F6"}, tag="Verified／Derived（Q4）")
qrow(6, "  基礎架構軟體", {"B": 27.029, "C": 6.796, "D": 7.178, "E": 8.752, "F": "=q4_sw"}, tag="Verified／Interested-party（Q4）")
qrow(7, "AI 半導體營收", {"B": "=co_ai26/(1+1.86)", "C": 8.4, "D": 10.8, "E": 16.7, "F": "=q4_ai"}, tag="Verified／Interested-party（Q4）／Derived（FY2025）", s2="FY2025＝公司 FY2026 展望 58 ÷（1＋186%）≈ 20.3（JPMorgan：約 20，二手）")
qrow(8, "  AI 網通（實際推估；markets/network rep_avgo26）", {"G": "=rep_avgo26"}, ann="", tag="Derived", s1=REPB[1], s2="Q3 FY26：XPU 占 AI 73%（法說，二手，單一來源）")
qrow(9, "  XPU（＝AI − 網通）", {"G": "=G7-G8"}, ann="", tag="Derived", s1="—", s2="FY2026 XPU 占 AI 營收（下一列）")
qrow(10, "  XPU ÷ AI", {"G": "=G9/G7"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="Q3 FY26 73%（法說，二手）")
qrow(11, "非 AI 半導體（＝半導體 − AI）", {c: f"={c}5-{c}7" for c in "BCDEF"}, tag="Derived")
qrow(12, "non-GAAP 毛利", {"B": 50.245, "C": 14.868, "D": 17.109, "E": 22.191, "F": "=F4*q4_gm"}, tag="Verified／Interested-party（Q4）")
qrow(13, "毛利率", {c: f"={c}12/{c}4" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="FY2025 78.6%；Q4 FY26 指引約 73%")
qrow(14, "non-GAAP 營業費用", {"B": 8.248, "C": 2.042, "D": 2.181, "E": 2.096, "F": "=F12-F15"}, tag="Verified／Derived（Q4＝毛利 − 營業利益）")
qrow(15, "non-GAAP 營業利益", {"B": 41.997, "C": 12.826, "D": 14.928, "E": 20.095, "F": "=F4*q4_oim"}, tag="Verified／Interested-party（Q4）")
qrow(16, "營業利益率", {c: f"={c}15/{c}4" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(17, "non-GAAP 利息費用", {"B": 3.072, "C": 0.746, "D": 0.745, "E": 0.703, "F": "=q4_int"}, tag="Verified／Derived（Q4）")
qrow(18, "non-GAAP 其他收益", {"B": 0.294, "C": 0.118, "D": 0.118, "E": 0.098, "F": "=q4_oth"}, tag="Verified／Derived（Q4）")
qrow(19, "稅前利益", {c: f"={c}15-{c}17+{c}18" for c in "BCDEF"}, tag="Derived")
qrow(20, "non-GAAP 所得稅", {"B": 5.491, "C": 2.013, "D": 2.227, "E": 3.118, "F": "=F19*q4_tx"}, tag="Verified／Derived（Q4）")
qrow(21, "有效稅率", {c: f"={c}20/{c}19" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(22, "non-GAAP 淨利", {c: f"={c}19-{c}20" for c in "BCDEF"}, tag="Derived")
qrow(23, "  公司揭露 non-GAAP 淨利", {"B": 33.728, "C": 10.185, "D": 12.074, "E": 16.372}, ann="", tag="Verified")
qrow(24, "non-GAAP 稀釋股數（十億股）", {"B": 4.943, "C": 4.957, "D": 4.940, "E": 4.937, "F": "=q4_sh"}, ann="=AVERAGE(C24:F24)", tag="Verified／Derived（Q4）")
qrow(25, "non-GAAP EPS", {c: f"={c}22/{c}24" for c in "BCDEF"}, fmt="0.000", ann="=G22/G24", tag="Derived", s1="淨利 ÷ 稀釋股數", s2="—")
qrow(26, "  公司揭露 non-GAAP EPS", {"B": 6.82, "C": 2.05, "D": 2.44, "E": 3.32}, fmt="0.00", ann="", tag="Verified")
qrow(27, "檢查：淨利 − 公司揭露（應約為 0；四捨五入）", {c: f"={c}22-{c}23" for c in "BCDE"}, fmt="0.000", ann="", tag="—", s1="—", s2="—")
qrow(28, "檢查：Q1＋Q2＋Q3 − 九個月合計（營收、毛利、營業利益、淨利；應為 0）", {"G": "=(C4+D4+E4-71.089)+(C12+D12+E12-54.168)+(C15+D15+E15-47.849)+(C22+D22+E22-38.631)"}, fmt="0.000", ann="", tag="—", s1=Q3PR + "（九個月欄）", s2="—")
Q.cell(30, 1, "FY2026 錨點與共識、公司展望的距離（參考；校準未使用共識）").font = BOLD
Q.cell(31, 1, "FY2026 營收 ÷ 共識 − 1"); Q["B31"] = "=G4/cons_rev26-1"
Q.cell(32, 1, "FY2026 EPS ÷ 共識 − 1"); Q["B32"] = "=G25/cons_eps26-1"
Q.cell(33, 1, "FY2026 AI 營收 ÷ 公司展望 58 − 1"); Q["B33"] = "=G7/co_ai26-1"
Q.cell(34, 1, "FY2026 營收年增"); Q["B34"] = "=G4/B4-1"
Q.cell(35, 1, "FY2026 營業費用年增"); Q["B35"] = "=G14/B14-1"
Q.cell(36, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B36"] = "=B35/B34"
for rr in range(31, 37): Q[f"B{rr}"].number_format = "0.0%" if rr < 36 else "0.00"
for nm, rr in (("rev", 4), ("semi", 5), ("sw", 6), ("ai", 7), ("net", 8), ("xpu", 9), ("nai", 11), ("gp", 12), ("opx", 14), ("oi", 15), ("int", 17), ("oth", 18), ("tax", 20), ("ni", 22), ("sh", 24), ("eps", 25)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEFG": Q.column_dimensions[c].width = 15
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 90; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)          # 有由下而上推導的年度（FY＝CY）；FY2026 用來校準
def uexp(k, y, s):
    if y <= 2027: return f"u_{k}_{y % 100}"
    if y == 2028: return f"u_{k}_28*um28_{s}"
    return f"u_{k}_28*(1+g29_{k})*um29_{s}"
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", uf="1", sht="0", sho="0", pfc="1", fh="f_hbm", mh="m_hbm", nf="1", gs27="g_sw27", gs="g_sw", opxa="0", flat29=False)
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.0", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    yu = lambda y: 2028 if (P["flat29"] and y == 2029) else y
    sec_("A. XPU（由下而上；D 欄＝FY2026 校準年，E–G＝FY2027–FY2029；FY(t)＝CY(t)）")
    for k in XP:
        row("u_" + k, f"{PNAME[k]} 市場顆數", "百萬顆", {c_(y): "=" + uexp(k, yu(y), s) for y in BU}, fmt="0.00")
    shx = {"tpu": P["sht"], "oth": P["sho"]}
    for k in XP:
        row("s_" + k, f"Broadcom 份額：{PNAME[k]}", "%", {c_(y): (f"=sh_{k}_{min(y, 2028) % 100}" if y == 2026 else f"=MIN(1,MAX(0,sh_{k}_{min(y, 2028) % 100}+{shx[k]}))") for y in BU}, fmt="0.0%")
    for k in XP:
        row("b_" + k, f"Broadcom 顆數：{PNAME[k]}", "百萬顆", {c_(y): f"={c_(y)}{R['u_' + k]}*{c_(y)}{R['s_' + k]}" + ("" if y == 2026 else f"*{P['uf']}") for y in BU}, fmt="0.00")
    for k in XP:
        row("g_" + k, f"每顆 HBM：{PNAME[k]}", "GB", {c_(y): f"=gb_{k}_{min(y, 2028) % 100}" for y in BU}, fmt="0")
    row("hpx", "HBM 單價", "US$/GB", {c_(y): f"=hbm_{s}_{y % 100}" for y in BU}, fmt="0.00")
    for k in XP:
        row("c_" + k, f"每顆核心售價：{PNAME[k]}", "US$K", {c_(y): f"=core_{k}" + ("" if y == 2026 else f"*{P['pfc']}") for y in BU}, fmt="0.00")
    for k in XP:
        row("h_" + k, f"每顆 HBM 金額（全部）：{PNAME[k]}", "US$K", {c_(y): f"={c_(y)}{R['g_' + k]}*{c_(y)}{R['hpx']}/1000" for y in BU}, fmt="0.00")
    for k in XP:
        row("a_" + k, f"每顆 Broadcom 售價（核心＋代購 HBM）：{PNAME[k]}", "US$K", {c_(y): f"={c_(y)}{R['c_' + k]}+{P['fh']}*{c_(y)}{R['h_' + k]}" for y in BU}, fmt="0.00",
            note="核心售價 ＋ 代購比例 × HBM 金額（未乘 k_x）" if k == "tpu" else "")
    row("xraw", "XPU 營收（未校準）", "US$B", {c_(y): "=" + "+".join(f"{c_(y)}{R['b_' + k]}*{c_(y)}{R['a_' + k]}" for k in XP) for y in BU})
    row("kx", "XPU 校準係數 k_x ＝ FY2026 錨點 XPU ÷ 未校準", "x", {"D": f"=a26_xpu/D{r - 1}"}, fmt="0.000", note="Derived：以已發生的實績（Q1–Q3＋Q4 指引）校準；FY2027 起沿用。顆數層、份額與核心單價無法分開識別，k_x 是合併係數")
    row("xpu", "XPU 營收", "US$B", {c_(y): f"={c_(y)}{R['xraw']}*$D${R['kx']}" for y in BU})
    row("xh", "  其中代購 HBM", "US$B", {c_(y): f"=$D${R['kx']}*{P['fh']}*(" + "+".join(f"{c_(y)}{R['b_' + k]}*{c_(y)}{R['h_' + k]}" for k in XP) + ")" for y in BU})
    row("xc", "  其中核心（設計＋邏輯晶片＋封裝）", "US$B", {c_(y): f"={c_(y)}{R['xpu']}-{c_(y)}{R['xh']}" for y in BU})
    row("xu", "Broadcom XPU 顆數（未乘 k_x）", "百萬顆", {c_(y): f"={c_(y)}{R['b_tpu']}+{c_(y)}{R['b_oth']}" for y in BU}, fmt="0.00")
    sec_("B. AI 網通、非 AI 半導體、軟體")
    row("nraw", "AI 網通（markets/network 歸屬 × 情境顆數係數）", "US$B", {c_(y): f"=net_{yu(y) % 100}*{umx(yu(y), s)}" + ("" if y == 2026 else f"*{P['nf']}") for y in BU})
    row("kn", "網通校準係數 k_n ＝ 實際推估 ÷ 歸屬（FY2026）", "x", {"D": f"=a26_net/D{r - 1}"}, fmt="0.000", note="Derived：markets/network 的 Broadcom 歸屬比例為 Assumed，2026 較實際推估低；FY2027 起沿用")
    row("net", "AI 網通營收", "US$B", {c_(y): f"={c_(y)}{R['nraw']}*$D${R['kn']}" for y in BU})
    row("ai", "AI 半導體營收", "US$B", {"C": "=a25_ai", **{c_(y): f"={c_(y)}{R['xpu']}+{c_(y)}{R['net']}" for y in BU}})
    row("xs", "XPU ÷ AI", "%", {c_(y): f"={c_(y)}{R['xpu']}/{c_(y)}{R['ai']}" for y in BU}, fmt="0.0%")
    row("nai", "非 AI 半導體", "US$B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+g_nai)" for y in BU[1:]}})
    row("sw", "基礎架構軟體", "US$B", {"C": "=a25_sw", "D": "=a26_sw", "E": f"=D{r}*(1+{P['gs27']})", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gs']})" for y in BU[2:]}})
    sec_("C. 報告值（FY2025A、FY2026 為實績與錨點；FY2027 起為由下而上）")
    row("rev", "營收", "US$B", {"C": "=a25_rev", "D": "=a26_rev", **{c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}+{c_(y)}{R['sw']}" for y in BU[1:]}})
    row("bur", "  由下而上營收（FY2026 應等於錨點）", "US$B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}+{c_(y)}{R['sw']}" for y in BU})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("gpo", "毛利：網通＋非 AI＋軟體＋代購 HBM", "US$B", {c_(y): f"=gm_net*{c_(y)}{R['net']}+gm_nai*{c_(y)}{R['nai']}+gm_sw*{c_(y)}{R['sw']}+{P['mh']}*{c_(y)}{R['xh']}" for y in BU})
    row("mc", "XPU 核心毛利率 m_core ＝（錨點毛利 − 其他毛利）÷ 核心營收（FY2026）", "x", {"D": f"=(a26_gp-D{r - 1})/D{R['xc']}"}, fmt="0.000", note="Derived：讓 FY2026 毛利率等於錨點；FY2027 起的毛利率變化來自營收組合（代購 HBM 隨 HBM 漲價與容量上升）")
    row("gp", "毛利", "US$B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"={c_(y)}{R['gpo']}+$D${R['mc']}*{c_(y)}{R['xc']}" for y in BU[1:]}})
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("hsh", "代購 HBM ÷ 營收", "%", {c_(y): f"={c_(y)}{R['xh']}/{c_(y)}{R['rev']}" for y in BU}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(g_opx_min,phi*{c_(y)}{R['g']}))" for y in BU[1:]}}, fmt="#,##0.00")
    row("opx", "營業費用", "US$B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}}, fmt="#,##0.00")
    row("oi", "營業利益", "US$B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("int", "利息費用", "US$B", {"C": "=a25_int", "D": "=a26_int", **{c_(y): "=int_f" for y in BU[1:]}}, fmt="0.00")
    row("oth", "其他收益", "US$B", {"C": "=a25_oth", "D": "=a26_oth", **{c_(y): "=oth_f" for y in BU[1:]}}, fmt="0.00")
    row("etr", "有效稅率", "%", {"C": "=a25_tax/(a25_oi-a25_int+a25_oth)", "D": "=a26_tax/(a26_oi-a26_int+a26_oth)", **{c_(y): "=tx_f" for y in BU[1:]}}, fmt="0.0%")
    row("ni", "淨利（non-GAAP）", "US$B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}-{c_(y)}{R['int']}+{c_(y)}{R['oth']})*(1-{c_(y)}{R['etr']})" for y in BU[1:]}})
    row("sh", "稀釋股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in BU[1:]}}, fmt="0.000")
    row("eps", "EPS（non-GAAP）", "US$", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh']}" for y in FYS}, fmt="0.00")
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL：三情境 ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Broadcom 損益（US$B；non-GAAP）。FY2025A 實績、FY2026 錨點（Q1–Q3 實績＋Q4 指引），FY2027 起為由下而上。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始；HBM 價格 2027 三情境相同）"
PB = {}; r0 = 70
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 半導體", "US$B", "#,##0.0"), ("xpu", "    XPU", "US$B", "#,##0.0"),
                                ("net", "    AI 網通", "US$B", "#,##0.0"), ("sw", "  基礎架構軟體", "US$B", "#,##0.0"),
                                ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利（non-GAAP）", "US$B", "#,##0.0"),
                                ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            if key in ("xpu", "net") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 由下而上營收 − 錨點；FY2026 由下而上 XPU − 錨點 XPU（k_x、k_n 校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['bur']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：XPU 差（US$B）"); PL[f"D{sr}"] = f"=D{PB['A']['xpu']}-a26_xpu"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 70, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：ASIC 每顆規格（權威）、金額份額（accel_share）、校準值 ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — ASIC 每顆規格（本專案權威，規則第 13 條）與加速器金額份額。顆數：百萬顆（日曆年；與 HBM、台積電、網通、NVIDIA 模型同一份）；每顆價：US$K"
A_ = PB["A"]
CYS = (2026, 2027, 2028, 2029); cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026（FY2026）", "CY2027（FY2027）", "CY2028（FY2028）", "CY2029（FY2029）", "說明"])
ur = 4
UN.cell(ur, 1, "1. ASIC 每顆規格（情境 A；Broadcom 營收口徑＝核心售價＋代購 HBM，乘 k_x）").font = BOLD; ur += 1
SPEC = {}
for k in XP:
    for key, lab, unit, f, fmt, note in (
        ("s", f"Broadcom 份額：{PNAME[k]}", "%", lambda y: f"=PL!{FC[y]}{A_['s_' + k]}", "0.0%", "In 頁 B 節（Assumed）"),
        ("b", f"Broadcom 顆數：{PNAME[k]}", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['b_' + k]}", "0.00", "市場層顆數 × 份額"),
        ("g", f"每顆 HBM：{PNAME[k]}", "GB", lambda y: f"=PL!{FC[y]}{A_['g_' + k]}", "0", "markets/memory UNITS"),
        ("h", f"每顆 HBM 金額：{PNAME[k]}", "US$K", lambda y: f"=PL!{FC[y]}{A_['h_' + k]}", "0.00", "容量 × HBM 單價（情境 A）"),
        ("p", f"每顆 Broadcom 營收（校準後）：{PNAME[k]}", "US$K", lambda y: f"=PL!{FC[y]}{A_['a_' + k]}*PL!$D${A_['kx']}", "0.00", "（核心＋代購比例 × HBM）× k_x"),
        ("q", f"每顆總價（核心 × k_x ＋ 全部 HBM；金額份額用）：{PNAME[k]}", "US$K", lambda y: f"=PL!{FC[y]}{A_['c_' + k]}*PL!$D${A_['kx']}+PL!{FC[y]}{A_['h_' + k]}", "0.00", "與 NVIDIA 模組價同口徑（含全部 HBM，不論誰採購）")):
        UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
        for y in CYS:
            UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
        UN.cell(ur, 7, note); SPEC[(key, k)] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "2. 校準值（只用 FY2026 已發生的實績與當季指引；不用共識、不用公司對 FY2027–28 的展望）").font = BOLD; ur += 1
CAL = {}
for key, lab, f, fmt, note in (("kx", "XPU 校準係數 k_x", f"=PL!D{A_['kx']}", "0.000", "FY2026 錨點 XPU ÷ 未校準由下而上；顆數、份額、核心單價的合併係數"),
                               ("kn", "網通校準係數 k_n", f"=PL!D{A_['kn']}", "0.000", "rep_avgo26 ÷ markets/network 歸屬"),
                               ("mc", "XPU 核心毛利率 m_core", f"=PL!D{A_['mc']}", "0.0%", "讓 FY2026 毛利率等於錨點的殘差（Derived）"),
                               ("x26", "FY2026 錨點 XPU", "=a26_xpu", "0.0", "AI 57.6 − 網通 17.6"),
                               ("xr26", "FY2026 未校準 XPU", f"=PL!D{A_['xraw']}", "0.0", "")):
    UN.cell(ur, 1, lab); UN[f"C{ur}"] = f; UN[f"C{ur}"].number_format = fmt; UN.cell(ur, 7, note); CAL[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "3. 加速器金額份額（accel_share；各平台每顆總價 × 顆數；ASIC 用本頁第 1 節的 CY2027 每顆總價，其他平台讀 NVIDIA 模型）").font = BOLD; ur += 1
PRICE = {"bw": "pm_bw", "rb": "pm_rb", "ru": "pm_ru", "hp": "pm_hp", "amd": "pa_amd", "trn": "pa_trn",
         "tpu": f"Units!$D${SPEC[('q', 'tpu')]}", "oth": f"Units!$D${SPEC[('q', 'oth')]}"}
ycols = {y: cyc(y) for y in CYS}
SHR = {}
for s in SC:
    ttl = {"A": "情境 A：基準", "B": "情境 B：2028 顆數 ×0.8、2029 ×0.85", "C": "情境 C：2028–29 顆數 ×0.6"}[s]
    units = {(p, y): uexp(p, y, s) for p in PLAT for y in CYS}
    ur, SHR[s] = AS.write_share_block(UN, ur, ttl, CYS, ycols, units, PRICE, PLAT, PNAME)
    UN.cell(ur - 1, 1, "  ASIC 合計（TPU＋Trainium＋其他）金額份額").font = BOLD
    for y in CYS:
        c = cyc(y)
        UN[f"{c}{ur - 1}"] = f"={c}{SHR[s][('ds', 'tpu')]}+{c}{SHR[s][('ds', 'trn')]}+{c}{SHR[s][('ds', 'oth')]}"; UN[f"{c}{ur - 1}"].number_format = "0.0%"
    SHR[s]["asic_d"] = ur - 1
    UN.cell(ur, 1, "  ASIC 合計顆數份額").font = BOLD
    for y in CYS:
        c = cyc(y)
        UN[f"{c}{ur}"] = f"={c}{SHR[s][('us', 'tpu')]}+{c}{SHR[s][('us', 'trn')]}+{c}{SHR[s][('us', 'oth')]}"; UN[f"{c}{ur}"].number_format = "0.0%"
    SHR[s]["asic_u"] = ur
    UN.cell(ur + 1, 1, "  Broadcom ÷ ASIC 金額（ASIC 總價口徑）").font = BOLD
    for y in CYS:
        c = cyc(y)
        UN[f"{c}{ur + 1}"] = (f"=(PL!{FC[y]}{PB[s]['b_tpu']}*{c}{SPEC[('q', 'tpu')]}+PL!{FC[y]}{PB[s]['b_oth']}*{c}{SPEC[('q', 'oth')]})"
                              f"/({c}{SHR[s][('d', 'tpu')]}+{c}{SHR[s][('d', 'trn')]}+{c}{SHR[s][('d', 'oth')]})")
        UN[f"{c}{ur + 1}"].number_format = "0.0%"
    SHR[s]["avgo_d"] = ur + 1
    UN.cell(ur + 2, 1, "  Broadcom ÷ ASIC 顆數").font = BOLD
    for y in CYS:
        c = cyc(y)
        UN[f"{c}{ur + 2}"] = f"=PL!{FC[y]}{PB[s]['xu']}/({c}{SHR[s][('u', 'tpu')]}+{c}{SHR[s][('u', 'trn')]}+{c}{SHR[s][('u', 'oth')]})"
        UN[f"{c}{ur + 2}"].number_format = "0.0%"
    SHR[s]["avgo_u"] = ur + 2
    ur += 4
UN.cell(ur, 1, "口徑：每顆價為「晶片＋HBM」的模組層級價（NVIDIA＝GPU 模組價；ASIC＝本頁第 1 節總價，含全部 HBM，不論由誰採購），不含系統加價與網通。ASIC 每顆價的金額份額以 CY2027 價套用各年（accel_share 每平台一個價格）。").font = BOLD
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 17
UN.column_dimensions["G"].width = 70

# ---------------- Val：P/E ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（美元；non-GAAP EPS）。結論用「現價 ÷ FY2027、FY2028 EPS（模型與共識）」。十年 P/E 規則值本版不適用：GAAP 落後 P/E 受無形資產攤銷扭曲（2016 年為負、2018 年 6.5x、2024 年 137x），前瞻 P/E 歷史待 Bloomberg（10-27 抓取單）"
V["A2"] = "循環位置：共用規則（tools/cycle_position.py classify；AVGO 的 FY(t) 對照日曆年 t，不平移），見頂年度讀設備市場模型；本版只記錄判定，倍數待前瞻 P/E 歷史"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-22 收盤）"); V["B6"] = "=px"; V["B6"].number_format = "0.00"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）")
V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）")
V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）")
V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "現價 ÷ 共識 FY2027 EPS"); V["B11"] = "=px/cons_eps27"; V["B11"].number_format = "0.0"
V.cell(12, 1, "近 12 月 non-GAAP P/E（Q4 FY25＋Q1–Q3 FY26 EPS）"); V["B12"] = "=px/(eps_q4_25+Q!C26+Q!D26+Q!E26)"; V["B12"].number_format = "0.0"
V.cell(13, 1, "近 12 月 GAAP P/E（StockAnalysis）"); V["B13"] = "=pe_ttm_gaap"; V["B13"].number_format = "0.0"
V.cell(14, 1, "前瞻 P/E（StockAnalysis；口徑未說明）"); V["B14"] = "=pe_fwd_sa"; V["B14"].number_format = "0.0"
V["E8"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"; V["E12"] = "Andy 先驗 B9a 答「前瞻 35x」：接近近 12 月 non-GAAP 倍數，而非前瞻倍數（見 Recon D 節）"
V.cell(16, 1, "2. 循環位置（共用規則，FY(t) 對照 CY t）——只記錄判定；倍數待前瞻 P/E 歷史").font = BOLD
hdr(V, 17, ["情境", "設備市場見頂年度（CY）", "FY2027 判定", "FY2028 判定", "FY2029 判定", "規則"])
for i, s in enumerate(SC, 18):
    V.cell(i, 1, s); V.cell(i, 2, f"=pk_{s}")
    for j, y in enumerate((2027, 2028, 2029)):
        V.cell(i, 3 + j, "=" + CP.excel_rule(str(y), f"pk_{s}"))
    V.cell(i, 6, "tools/cycle_position.py classify（不平移）")
V.cell(22, 1, "3. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（情境 A 的 FY2028 EPS；約 2027-10 的價格看下一年 EPS）").font = BOLD
hdr(V, 23, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，C）", "對照：Andy 合理 P/E", "對照：現價 ÷ 共識 FY2027 EPS"])
vr = 24
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    V.cell(vr, 2, f"=px*A{vr}/PL!F{SUMR[('A', 'eps')]}").number_format = "0.0"
    V.cell(vr, 3, f"=px*A{vr}/PL!F{SUMR[('C', 'eps')]}").number_format = "0.0"
    V.cell(vr, 4, "=pe_fair_andy").number_format = "0.0"; V.cell(vr, 5, "=px/cons_eps27").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數愈低，現價愈便宜；倍數是否合理，待前瞻 P/E 十年歷史（Bloomberg）。")
vr += 2
V.cell(vr, 1, "4. 對照：以 Andy 合理 P/E 25x 計算的隱含價（只作對照，非規則值）").font = BOLD; vr += 1
hdr(V, vr, ["情境", "FY2027 EPS", "× 25x", "相對現價", "FY2028 EPS", "× 25x", "相對現價"]); vr += 1
IMP = {}
for s in SC:
    V.cell(vr, 1, s)
    V.cell(vr, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 3, f"=B{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 4, f"=C{vr}/px-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 6, f"=E{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
    IMP[s] = vr; vr += 1
vr += 1
V.cell(vr, 1, "5. 十年 GAAP P/E（Macrotrends，10 月 31 日；不適用——GAAP EPS 受攤銷扭曲，只作紀錄）").font = BOLD; vr += 1
hdr(V, vr, ["年底", "P/E（GAAP 落後）", "說明"]); vr += 1
for y, v in ((2016, None), (2017, 49.57), (2018, 6.50), (2019, 38.87), (2020, 49.77), (2021, 32.75), (2022, 16.88), (2023, 24.95), (2024, 137.36), (2025, None)):
    V.cell(vr, 1, str(y))
    if v is not None: V.cell(vr, 2, v).font = BLUE; V.cell(vr, 2).number_format = "0.00"
    V.cell(vr, 3, "GAAP EPS 為負，P/E 無意義" if y == 2016 else ("頁面資料只到 2025-08：未找到" if y == 2025 else "")); vr += 1
V.cell(vr, 1, "來源：macrotrends.net/stocks/charts/AVGO/broadcom/pe-ratio（Interested-party，資料商）；第二來源：未找到。前瞻 P/E 歷史待 Bloomberg 抓取單（10-27）").font = Font(italic=True)
V.column_dimensions["A"].width = 58
for c in "BCDEFG": V.column_dimensions[c].width = 18
V.column_dimensions["E"].width = 40

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；k_x、k_n、m_core 在每個案例內依 FY2026 實績重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("XPU 顆數係數 ×0.9", "FY2027 起 Broadcom XPU 顆數 ×0.9", dict(uf="uk_lo")), ("XPU 顆數係數 ×1.1", "×1.1", dict(uf="uk_hi")),
         ("TPU 份額 −15pt", "FY2027 起 Broadcom 占 TPU 顆數 −15 個百分點（MediaTek 版放量）", dict(sht="(-0.15)")), ("TPU 份額 +10pt", "+10 個百分點（上限 100%）", dict(sht="0.10")),
         ("其他 ASIC 份額 −15pt", "FY2027 起 −15 個百分點", dict(sho="(-0.15)")), ("其他 ASIC 份額 +15pt", "+15 個百分點", dict(sho="0.15")),
         ("核心售價 ×0.8", "FY2027 起 XPU 核心售價 ×0.8", dict(pfc="pfc_lo")), ("核心售價 ×1.2", "×1.2", dict(pfc="pfc_hi")),
         ("HBM 代購比例 0.3", "代購 30%（k_x、m_core 重算）", dict(fh="f_hbm_lo")), ("HBM 代購比例 1.0", "全部代購", dict(fh="f_hbm_hi")),
         ("代購 HBM 毛利率 0%", "零加成", dict(mh="m_hbm_lo")), ("代購 HBM 毛利率 20%", "20%", dict(mh="m_hbm_hi")),
         ("AI 網通 ×0.8", "FY2027 起網通 ×0.8", dict(nf="0.8")), ("AI 網通 ×1.2", "×1.2", dict(nf="1.2")),
         ("軟體成長 低", "FY2027 +8%、之後 +4%", dict(gs27="g_sw27_lo", gs="g_sw_lo")), ("軟體成長 高", "FY2027 +18%、之後 +12%", dict(gs27="g_sw27_hi", gs="g_sw_hi")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "ai", 2027, "#,##0.0"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "ai", 2028, "#,##0.0"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85；記憶體 B 路徑", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6；記憶體 C 路徑", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數情境 B、C 的差異從 FY2028 起。k_x 隨案例重算，所以 HBM 代購比例只改變營收組合與毛利率，不改 FY2026 水準。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司展望、共識、市場層。只作事後對照，不為對上數字而調參數。公司 FY2027、FY2028 AI 營收展望（115、230）為 Interested-party，不用於校準"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. AI 營收 vs 公司展望（Interested-party；校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 AI 營收 vs 公司展望 58", "=co_ai26", pl("ai", 2026), "公司（法說）", "FY2026 為錨點（Q1–Q3 實績＋Q4 指引 57.6）")
rec("FY2027 AI 營收 vs 公司展望 115", "=co_ai27", pl("ai", 2027), "公司（法說）", "")
rec("FY2027 AI 營收 vs JPMorgan 150", "=jpm_ai27", pl("ai", 2027), "JPMorgan（KuCoin 轉述）", "")
rec("FY2028 AI 營收 vs 公司展望 230", "=co_ai28", pl("ai", 2028), "公司（法說）", "差距主要在顆數：公司說法含 Anthropic 2028 年再加約 10 GW、OpenAI 5 GW 以上")
RC.cell(rr_, 1, "B. What must be true：要達到公司展望，Broadcom XPU 顆數需要是模型的幾倍（網通、核心售價、HBM 不變）").font = BOLD; rr_ += 1
rec("FY2027：所需 XPU 顆數倍數", None, lambda s: f"=(co_ai27-PL!E{PB[s]['net']})/PL!E{PB[s]['xpu']}", "—", "倍數 >1＝公司展望需要比市場層更多的 Broadcom 顆數", fmt="0.00")
rec("FY2028：所需 XPU 顆數倍數", None, lambda s: f"=(co_ai28-PL!F{PB[s]['net']})/PL!F{PB[s]['xpu']}", "—", "", fmt="0.00")
rec("FY2028：所需 Broadcom XPU 顆數（百萬顆）vs 模型", f"=PL!F{PB['A']['xu']}", lambda s: f"=(co_ai28-PL!F{PB[s]['net']})/PL!F{PB[s]['xpu']}*PL!F{PB[s]['xu']}", "—", "對照值＝模型 A 的 Broadcom XPU 顆數", fmt="0.00", gap=False)
RC.cell(rr_, 1, "C. 共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "StockAnalysis", "錨點")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "StockAnalysis", "錨點（Q4 依指引與 In 頁假設）", fmt="0.00")
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "StockAnalysis", "")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "StockAnalysis", "", fmt="0.00")
rec("FY2027 共識隱含 AI 營收（共識營收 − 模型非 AI − 模型軟體）", None, lambda s: f"=cons_rev27-PL!E{PB[s]['nai']}-PL!E{PB[s]['sw']}", "推導", "與公司展望 115 比較", gap=False)
RC.cell(rr_, 1, "D. 份額與顆數（先驗比對的機制檢查）").font = BOLD; rr_ += 1
rec("2027 年 ASIC 金額份額（模組層級）vs 先驗比對推導 14–19% 中點", 0.165, lambda s: f"=Units!D{SHR[s]['asic_d']}", "先驗比對_AVGO（Drive）", "", fmt="0.0%", pt=True)
rec("2027 年 ASIC 顆數份額 vs JPMorgan 53.6%", "=jpm_u27", lambda s: f"=Units!D{SHR[s]['asic_u']}", "JPMorgan（二手）；市場層同一來源", "非獨立檢驗（市場層顆數即取自 JPMorgan）", fmt="0.0%", pt=True)
rec("2027 年 Broadcom ÷ ASIC 顆數 vs Counterpoint 60%", "=cp_sh27", lambda s: f"=Units!D{SHR[s]['avgo_u']}", "Counterpoint（二手；口徑未說明）", "", fmt="0.0%", pt=True)
rec("2027 年 Broadcom ÷ ASIC 金額 vs Counterpoint 60%", "=cp_sh27", lambda s: f"=Units!D{SHR[s]['avgo_d']}", "同上", "JPMorgan：高階 ASIC 80–85%", fmt="0.0%", pt=True)
rec("2027 年 NVIDIA 金額份額", None, lambda s: f"=Units!D{SHR[s][('ds', 'nv')]}", "accel_share", "NVIDIA 模型同一套價格", fmt="0.0%")
rec("每顆價差：NVIDIA Rubin 模組 ÷ TPU 總價（CY2027）", None, lambda s: f"=pm_rb/Units!D{SPEC[('q', 'tpu')]}", "—", "顆數份額約為金額份額數倍的原因", fmt="0.0")
rec("近 12 月 non-GAAP P/E vs Andy 先驗「前瞻 35x」", 35, lambda s: "=Val!B12", "先驗_Andy_AVGO B9a（Drive）", "35x 接近近 12 月倍數；前瞻（FY2027 EPS）倍數見 Val 頁", fmt="0.0")
RC.cell(rr_, 1, "E. 校準與網通").font = BOLD; rr_ += 1
rec("k_x（XPU 校準係數）", None, lambda s: f"=PL!D{PB[s]['kx']}", "FY2026 錨點 ÷ 未校準", "接近 1＝市場層顆數 × 份額 × 假設單價與實績一致", fmt="0.000")
rec("k_n（網通校準係數）", None, lambda s: f"=PL!D{PB[s]['kn']}", "rep_avgo26 ÷ markets/network 歸屬", "markets/network 的 Broadcom 歸屬比例偏低（Assumed）", fmt="0.000")
rec("m_core（XPU 核心毛利率，Derived）", None, lambda s: f"=PL!D{PB[s]['mc']}", "FY2026 毛利殘差", "高＝設計與邏輯晶片加價高；代購 HBM 稀釋整體", fmt="0.0%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "公司 Q4 FY26 指引約 73%，續降", fmt="0.0%", pt=True)
rec("FY2027 毛利率 vs 共識隱含（約 72%：EPS 19.40 × 股數 ÷（1−稅）＋利息＋營業費用，÷ 共識營收）", "=(cons_eps27*a26_sh/(1-tx_f)+int_f-oth_f+PL!E" + str(PB['A']['opx']) + ")/cons_rev27", pl("gm", 2027), "推導（共識無毛利率）", "", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "F. 客戶融資曝險（資訊；不進基準 EPS）").font = BOLD; rr_ += 1
rec("Anthropic 放款＋OpenAI 融資安排（標題值）vs FY2027 AI 營收（模型）", "=fin_anth+fin_oai", pl("ai", 2027), "top1markets 引 Reuters、WSJ、Bloomberg", "媒體：約 92B ≈ 公司 FY2027 AI 目標的 80%；非資產負債表曝險，部分在銀行團；另有一家客戶機架租賃擔保約 29B（季報），可能與上列重疊，不加總", gap=False)
rec("兩家合計 ÷ FY2027 AI 營收（模型）", None, lambda s: f"=(fin_anth+fin_oai)/PL!E{PB[s]['ai']}", "—", "", fmt="0%")
rec("兩家合計 ÷ FY2027 淨利（模型）", None, lambda s: f"=(fin_anth+fin_oai)/PL!E{PB[s]['ni']}", "—", "若全數損失的量級（不含回收與銀行團分擔）", fmt="0.0")
RC.cell(rr_, 1, "G. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 Broadcom AI 營收 ÷ 部署 GW 對照 Tokenomics IF_NonNV_Hold_TPU 等比例"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 80
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 40; RC.column_dimensions["J"].width = 70

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝Broadcom 8-K 新聞稿，已逐份抓取；法說逐字稿只經二手轉述）"; SO["A1"].font = BOLD
srcs = [("公司：FY2025 與 Q4 FY25 新聞稿（8-K）", EDG + "000173016825000116/avgo-11022025x8kxex99.htm"),
        ("公司：Q1 FY26 新聞稿（8-K）", EDG + "000173016826000011/avgo-02012026x8kxex99.htm"),
        ("公司：Q2 FY26 新聞稿（8-K）", EDG + "000173016826000051/avgo-05032026x8kxex99.htm"),
        ("公司：Q3 FY26 新聞稿（含 Q4 指引、九個月合計）", Q3PR),
        ("二手：Q3 FY26 法說逐字稿（XPU 73%、客戶、毛利率路徑、FY2027–28 展望）", "https://www.investing.com/news/transcripts/earnings-call-transcript-broadcom-tops-q3-2026-estimates-as-ai-sales-surge-93CH-4886849"),
        ("二手：BigGo Finance 法說摘要", "https://finance.biggo.com/news/US_AVGO_2026-09-02"),
        ("二手：StockAnalysis 共識、現價、前瞻 P/E", "https://stockanalysis.com/stocks/avgo/forecast/"),
        ("二手：GuruFocus 前瞻 P/E", "https://www.gurufocus.com/term/forward-pe-ratio/AVGO"),
        ("二手：valueinvesting.io 前瞻 P/E 49.62（口徑不明，不採）", "https://valueinvesting.io/AVGO/metric/forward-pe"),
        ("二手：top1markets 客戶融資與股價（2026-10-22）", "https://www.top1markets.com/news/broadcom-stock-ai-revenue-230-billion-financing-risk"),
        ("二手：JPMorgan（KuCoin 轉述）：ASIC 顆數、Broadcom AI 營收", "https://www.kucoin.com/news/flash/jpmorgan-ai-custom-chip-shipments-to-surpass-gpus-by-2027"),
        ("二手：Counterpoint（Yahoo 轉述）：Broadcom 2027 ASIC 份額約 60%", "https://finance.yahoo.com/news/broadcom-set-dominate-custom-ai-163116560.html"),
        ("二手：Macrotrends GAAP P/E", "https://www.macrotrends.net/stocks/charts/AVGO/broadcom/pe-ratio"),
        ("市場層：markets/memory/market.py（UNITS）；tsmc、skhynix、markets/network、nvidia 成品（本 repo，唯讀）", ""),
        ("待補（人工）：前瞻 P/E 十年歷史（Bloomberg，10-27）；各客戶營收；XPU 產品別毛利率；HBM 代購比例", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 80; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Broadcom（AVGO）收支與估值模型 v0.1（AI 半導體 Project：運算層第二個樣板，ASIC＋網通）　2026-10-23", True),
 ("版本：v0.1（初版）。ASIC 每顆規格（Broadcom 份額、核心售價、HBM 代購）在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）。", False),
 ("模型命題：Broadcom 的 AI 營收＝各 XPU 平台的顆數 × Broadcom 份額 × 每顆售價（核心＋代購 HBM）＋網通附掛；2027–28 年的增量集中在 Anthropic、OpenAI 兩家靠融資部署的客戶。股價約 17 倍 FY2027 EPS：這個折價是在定價客戶信用風險，還是低估了營收？", False),
 ("核心問題：以市場層顆數往前推，Broadcom 的 AI 營收與 EPS 能否支撐公司展望（FY2027 約 115B、FY2028 約 230B）；差距要靠多少額外顆數補上。", False), ("", False),
 ("會計年度：FY(t) 於 t 年 10 月底／11 月初結束，10/12 落在日曆年 t，FY(t) 直接對照 CY(t)，不做平移。預測期 FY2027–FY2029（三年）。口徑：公司 non-GAAP（排除股票薪酬、攤銷、重組）。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("XPU：Google TPU 與「其他 ASIC」顆數（markets/memory UNITS × 情境係數）× Broadcom 份額（Assumed）×（核心售價＋代購比例 × 每顆 HBM 容量 × HBM 單價）× k_x（FY2026 實績校準）→ XPU 營收。", False),
 ("AI 網通：markets/network 的「Broadcom 網通（本模型歸屬）」× 情境顆數係數 × k_n（FY2026 實際推估校準）。非 AI 半導體、軟體：FY2026 錨點 × Assumed 成長率。", False),
 ("毛利率：網通、非 AI、軟體、代購 HBM 各用 Assumed 毛利率；XPU 核心毛利率 m_core 為 FY2026 殘差（Derived）。FY2027 起毛利率下滑來自代購 HBM 占比上升（HBM 2027 年 +80%、每顆容量上升）。", False),
 ("校準：FY2026 ＝ Q1–Q3 實績＋Q4 指引（營收 34.8、AI 21.7、營業利益率 66%）。不用共識、不用公司對 FY2027–28 的展望（只在 Recon 對帳）。", False),
 ("估值：結論用現價 ÷ FY2027、FY2028 EPS；十年 GAAP P/E 規則值不適用（攤銷扭曲），前瞻 P/E 歷史待 Bloomberg。循環位置（共用規則，不平移）只作紀錄：設備市場見頂年度 " + "、".join(f"{s_} {PEAK[s_]}" for s_ in SC) + "。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文＝SEC EDGAR 8-K；法說內容只經二手轉述（標二手）。", False),
 ("已知限制：(1) k_x 是顆數、份額、核心單價的合併係數，無法分開識別；(2) Broadcom 各平台份額、核心售價、HBM 代購比例皆為 Assumed，FY2027 起的毛利率對代購比例高度敏感（Sens）；(3) 「其他 ASIC」含昇騰、Maia 等非 Broadcom 晶片，份額為粗估；(4) 客戶融資曝險不進基準 EPS（Recon F）；(5) 十年倍數待補；(6) R4 產業部署路徑未建（Recon 預留列）。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("AVGO_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "SHR": {s: {str(k): v for k, v in R.items()} for s, R in SHR.items()},
               "SPEC": {f"{k[0]}|{k[1]}": v for k, v in SPEC.items()}, "CAL": CAL, "CHK": CHK_ROW, "SENSN": SENS_N, "IMP": IMP, "SB": SB},
              open(os.environ["AVGO_ROWMAP"], "w"), ensure_ascii=False, indent=1)
