# 東京精密（Tokyo Seimitsu／ACCRETECH，7729.T）收支與估值模型 v0.1.1（AI 半導體 Project：設備與材料）
# 用法：python3 accretech/build_accretech.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度 4 月到隔年 3 月：FY2025＝2025-04～2026-03（公司「2026年3月期」）。金額：十億日圓（¥B）；EPS：日圓。
# 結構比照 disco/build_disco.py：README、In、Q、Hist、PL、Val、Recon、Sens、Sources。
# 設備支出讀 markets/equipment 成品 Mkt 頁的 A／B／C 同字母路徑；循環位置用 tools/cycle_position.py。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); FADE = PatternFill("solid", fgColor="E7E6E6")
wb = openpyxl.Workbook()
SC = ("A", "B", "C")
SCN = {"A": "順利", "B": "平平", "C": "不如預期"}
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- 讀入上游成品（唯讀；不修改） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
EQ_F = only("markets/equipment/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ws_ = openpyxl.load_workbook(EQ_F, data_only=True)["Mkt"]
EQ = {}       # EQ[(情境, 項目)] = ([CY2025..CY2029], 來源)
blk = None
for r in ws_.iter_rows():
    a = r[0].value
    if isinstance(a, str) and a[:3] in ("5A.", "5B.", "5C."):
        blk = a[1]; continue
    if blk and isinstance(a, str):
        for key, lab in (("tst", "測試與其他封裝"), ("apk", "先進封裝〔AI〕"), ("wfe", "晶圓前段設備"), ("tot", "設備總額（情境")):
            if a.startswith(lab) and (blk, key) not in EQ:
                EQ[(blk, key)] = ([r[2 + j].value for j in range(5)], f"{rel(EQ_F)}｜Mkt!C{r[0].row}:G{r[0].row}")
assert len(EQ) == 12 and all(v is not None for k in EQ for v in EQ[k][0]), sorted(EQ)
PEAK, PEAK_SRC = CP.peak_years(ROOT)

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億日圓（¥B）；EPS：日圓。每個輸入附標記與來源；找不到者寫「未找到」，公司未揭露者寫「公司不揭露」"
IN["A2"] = "標記：Verified（公司文件、交易所資料）／Interested-party（寫明是誰、立場）／Analogy（寫明對照與差異）／Assumed（一律給低／高區間）／Derived（寫明換算）。SemiAnalysis 未使用。公司原文（accretech.com、JPX 適時開示）以網頁擷取工具讀取；表格錯位處以算術核對"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
TS25 = "公司 2026年3月期 決算短信（JPX 適時開示，2026-05-13）"
FS25 = "公司 2026年3月期 決算說明補足資料（FS_2026_4q_J_rev，2026-05）"
TSQ1 = "公司 2027年3月期 第1四半期 決算短信（JPX，2026-08-04）"
BRQ1 = "公司 2026年度第1四半期 決算說明資料（BR_2027-3_1q_J_2，2026-08-04）"
QAQ1 = "公司 2026年度第1四半期 決算說明會 Q&A（QA_2027-3_1q_J，2026-08-04）"
MATSUI = "松井證券 決算情報（finance.matsui.co.jp/stock/7729/settlement）"

sec("A. FY2025 實績（2025-04～2026-03；校準錨）")
inp("a25_rev", "營收 FY2025", "¥B", None, 166.839, None, "Verified", TS25 + "：166,839 百萬日圓（+10.8%）", MATSUI + "：166,839")
inp("a25_gp", "毛利 FY2025", "¥B", None, 68.860, None, "Verified", FS25 + "：68,860（毛利率 41.3%）", "FY2024 62,453（41.5%）")
inp("a25_sga", "銷管費 FY2025", "¥B", None, 35.122, None, "Verified", FS25 + "：35,122", "毛利 − 銷管費 ＝ 33,738 ＝ 營業利益（算術核對）")
inp("a25_oi", "營業利益 FY2025", "¥B", None, 33.738, None, "Verified", TS25 + "：33,738（+13.6%）", MATSUI + "：33,738")
inp("a25_ni", "淨利（歸屬母公司）FY2025", "¥B", None, 24.739, None, "Verified", TS25 + "：24,739（−3.5%；FY2024 含較多營業外與特別利益，淨利 25,637）", MATSUI + "：24,739")
inp("a25_eps", "基本 EPS FY2025", "¥", None, 610.02, None, "Verified", TS25 + "：基本 610.02（稀釋 606.14）", MATSUI + "：610.02")
inp("spe_25", "半導體製造裝置（SPE）營收 FY2025", "¥B", None, 127.878, None, "Verified", FS25 + "：127,878（FY2024 113,481）", BRQ1 + "：四季 235／359／274／411 億（合計 1,279）")
inp("spe_op25", "SPE 營業利益 FY2025", "¥B", None, 28.404, None, "Verified", FS25 + "：28,404（22.2%）", BRQ1 + "：四季 40／83／53／107 億")
inp("met_25", "計測機器營收 FY2025", "¥B", None, 38.960, None, "Verified", FS25 + "：38,960（FY2024 37,053）", BRQ1 + "：四季 73／103／85／128 億")
inp("met_op25", "計測機器營業利益 FY2025", "¥B", None, 5.333, None, "Verified", FS25 + "：5,333（13.7%）", BRQ1 + "：四季 5／18／9／21 億")
inp("ord_25", "受注 FY2025（SPE＋計測）", "¥B", None, 163.096, None, "Verified", FS25 + "：163,096（SPE 123,396、計測 39,700）；期末受注殘 79,359", "只作參考（未進推導）")

sec("B. FY2026 錨點：公司指引（2026-08-04 修正，Interested-party）與第一季實績")
inp("g26_rev", "公司指引：FY2026 營收（2026-04～2027-03）", "¥B", None, 189.0, None, "Interested-party", TSQ1 + "：189,000（+13.3%）；5 月原指引 181,500", BRQ1 + "：上修理由＝受注殘與強勁詢價，主要上修下半；營收只含已確保材料的部分（" + QAQ1 + "）")
inp("g26_met", "公司指引：FY2026 計測機器營收", "¥B", None, "=40.0+1.5", None, "Interested-party", FS25 + "：5 月原指引計測 400 億", BRQ1 + "：計測通期 +15 億上修（下半為主）")
inp("g26_spe", "公司指引：FY2026 SPE 營收（＝營收 − 計測）", "¥B", None, "=g26_rev-g26_met", None, "Derived", "189.0 − 41.5 ＝ 147.5", "5 月原指引 SPE 141.5 ＋ 上修 6.0（" + BRQ1 + "：SPE +60 億）＝ 147.5，相同")
inp("g26_oi", "公司指引：FY2026 營業利益", "¥B", None, 44.0, None, "Interested-party", TSQ1 + "：44,000（+30.4%）；5 月原指引 40,000", "上半 19,100、下半 24,900（Derived）")
inp("g26_ni", "公司指引：FY2026 淨利", "¥B", None, 30.8, None, "Interested-party", TSQ1 + "：30,800（+24.5%）；EPS 758.83", "5 月原指引 28,000／689.87")
inp("g26_eps", "公司指引：FY2026 EPS", "¥", None, 758.83, None, "Interested-party", TSQ1 + "：758.83", "minkabu 轉載 758.81")
inp("m_spe26", "FY2026 SPE 營業利益率（公司說法）", "%", 0.25, 0.255, 0.26, "Interested-party", QAQ1 + "：通期半導體營業利益率假設「25-6%」（原文如此，應為 25–26%）", "計測營業利益＝指引營業利益 − SPE（Derived）；基準取中點")
inp("g_anc", "FY2026 實績相對公司指引（營收；兩部門同比例）", "%", -0.03, 0.0, 0.05, "Assumed", "基準＝公司指引（第二季尚未公布）；高：公司稱受注約為產能 1.5 倍、營收只含已確保材料，材料增加可再多數百億日圓（" + QAQ1 + "）", "低：營收認列遞延（交期 4–8 個月）。營業利益隨營收以邊際利潤率（v_s、v_m）變動")
inp("q1_rev", "FY2026 第一季營收（2026-04～06）", "¥B", None, 36.716, None, "Verified", TSQ1 + "：36,716（+18.9%）", MATSUI + "：36,716")
inp("q1_oi", "FY2026 第一季營業利益", "¥B", None, 5.920, None, "Verified", TSQ1 + "：5,920（+29.2%）", MATSUI)
inp("q1_ni", "FY2026 第一季淨利", "¥B", None, 4.667, None, "Verified", TSQ1 + "：4,667；EPS 115.01", MATSUI)
inp("q1_ord", "FY2026 第一季受注（全社）", "¥B", None, 65.8, None, "Interested-party", BRQ1 + "：658 億（+83%，SPE 532、計測 126；創新高；受注約一半來自中國）", "第一季末受注殘約 1,085 億（說明資料表格錯位，參考值）")
inp("gh1_rev", "公司指引：FY2026 上半營收", "¥B", None, 88.5, None, "Interested-party", TSQ1 + "：88,500（營業利益 19,100、淨利 13,400）", "第二季隱含營收 51.8（Derived）")
inp("shares", "加權平均股數（扣除員工持股信託）", "百萬股", None, 40.588440, None, "Verified", TSQ1 + "：第一季期中平均 40,588,440 股", "FY2025 平均 40,554,349（" + TS25 + "）")

sec("C. SPE 產品組合與客戶結構（公司說法；Interested-party）")
inp("w_pr", "探針台（prober）占 SPE 營收", "%", 0.70, 0.72, 0.74, "Interested-party", BRQ1 + "：FY2026 上半 SPE 營收構成 探針台「7 割前半」、加工裝置（切割、研削、CMP 等）「2 割後半」", "公司不揭露金額；FY2025 構成比未找到，假設與 FY2026 上半相同（Assumed）")
inp("w_cmp", "加工裝置中屬晶圓前段（CMP 等）的比例", "%", 0.10, 0.20, 0.35, "Assumed", "未找到：公司不揭露切割機、研削機、CMP 各自金額", "前段部分跟 markets/equipment「晶圓前段設備」，其餘跟「先進封裝〔AI〕」")
inp("hpc", "HPC（含生成 AI）占 SPE 營收（FY2026 第一季）", "%", None, 0.45, None, "Interested-party", QAQ1 + "：第一季 HPC 占 SPE 營收「4 割半ば」，其中邏輯 2／3、HBM 1／3；受注約 4 割", "FY2025 第一季約 25%（QA_2026_1q_J，2025-08-04）。只作說明，未進推導（上游設備路徑已含 AI）")
inp("chn", "中國占 SPE 受注（FY2026 第一季）", "%", 0.40, 0.45, 0.50, "Interested-party", BRQ1 + "：第一季受注「半分弱」來自中國", "營收占比未找到（補足資料表格錯位）；以受注比代替")

sec("D. 設備市場（markets/equipment 成品 Mkt 頁 5A／5B／5C，同字母情境；唯讀）")
for s in SC:
    for key, nm in (("tst", "測試與其他封裝設備"), ("apk", "先進封裝設備〔AI〕"), ("wfe", "晶圓前段設備")):
        for j, y in enumerate((25, 26, 27, 28, 29)):
            inp(f"{key}{s}_{y}", f"{nm} CY20{y}（情境 {s}）", "US$B", None, EQ[(s, key)][0][j], None, "Derived", EQ[(s, key)][1] if j == 0 else "同上", "markets/equipment v0.1.3：B、C 自 2028 起與 A 不同（設備市場模型的情境係數為 Assumed）" if (j == 0 and key == "tst" and s == "A") else "")
inp("g_tail", "CY2030 設備市場相對 CY2029（三情境相同；市場層只到 2029）", "%", -0.10, 0.0, 0.10, "Assumed", "無：FY2029＝3/4 × CY2029 ＋ 1/4 × CY2030，需要 CY2030；基準持平", "Sens 列 ±10%")

sec("E. 公司特有情境變數（方向一致：A 最有利、C 最不利）")
for s, b_, lo, hi, note in (("A", 0.03, 0.01, 0.05, "高溫控精度探針台（HBM、AI 邏輯）組合提升；探針台產能 FY2027 初約 1.2 倍、飯能新廠（" + QAQ1 + "）"),
                            ("B", 0.0, -0.01, 0.01, "與市場同步"),
                            ("C", -0.03, -0.05, -0.01, "中國受注占比高（約一半）、HBM 世代交替遞延（公司 2025-08 曾稱 HBM 探針台訂單延後一季）")):
    inp(f"exc_{s}", f"情境 {s}：SPE 相對設備市場的年超額成長（份額＋產品組合；FY2027 起）", "%/年", lo, b_, hi, "Assumed", note, "份額水準未找到（探針台、切割機均無一手市場調查）；只以變化進入模型")
for s, b_, lo, hi in (("A", 0.0, 0.0, 0.0), ("B", -0.05, -0.10, 0.0), ("C", -0.20, -0.30, -0.10)):
    inp(f"chh_{s}", f"情境 {s}：中國 SPE 營收水準變化（FY2027 一半、FY2028 起全額）", "%", lo if s != "A" else None, b_, hi if s != "A" else None, "Assumed", "中國高階邏輯需求「高水準持續」為公司前提（" + BRQ1 + "）；出口管制未在 Q&A 討論" if s == "A" else "C：中國需求消化或出口管制擴大；B：小幅回落", "SPE 營收 ×（1 ＋ 中國占比 × 本變數）")
for s, b_, lo, hi in (("A", 0.06, 0.04, 0.08), ("B", 0.04, 0.02, 0.06), ("C", 0.01, -0.02, 0.03)):
    inp(f"gmet_{s}", f"情境 {s}：計測機器營收年增（FY2027 起）", "%", lo, b_, hi, "Analogy", "公司歷史（Hist 頁）：計測營收 FY2018→FY2025 年複合約 2.7%、FY2020→FY2025 約 9.0%（FY2020 為疫情低點）", "與現在的差異：公司稱航太、國防、AI 資料中心與補助金帶動詢價（" + BRQ1 + "）；車用（HV）投資屬一次性" if s == "A" else "同上")
inp("madjC", "情境 C：FY2028 起 SPE 營業利益率額外調整（價格、稼動）", "pt", -0.02, -0.01, 0.0, "Assumed", "無", "")

sec("F. 利潤率（由 Hist 頁部門實績迴歸推得，不寫死）")
inp("v_s", "SPE 邊際營業利益率（ΔSPE 營業利益 ÷ Δ營收）", "x", None, "=SLOPE(Hist!J9:J16,Hist!I9:I16)", None, "Derived", "Hist 頁 FY2018–FY2025 八年 SPE 營業利益對營收最小平方法", "FY2027 起：SPE 營業利益＝FY2026 錨點 ＋ v_s ×（營收 − FY2026 營收）")
inp("v_m", "計測機器邊際營業利益率", "x", None, "=SLOPE(Hist!L9:L16,Hist!K9:K16)", None, "Derived", "Hist 頁 FY2018–FY2025 八年", "")
inp("v_s2", "SPE 邊際營業利益率（只用 FY2021–FY2025，Sens 用）", "x", None, "=SLOPE(Hist!J12:J16,Hist!I12:I16)", None, "Derived", "Hist 頁五年", "")
inp("nioi", "淨利 ÷ 營業利益（FY2027 起）", "x", 0.68, "=g26_ni/g26_oi", "=a25_ni/a25_oi", "Derived", "基準＝公司 FY2026 指引隱含 0.700（30.8 ÷ 44.0）；高＝FY2025 實績 0.733", "低 0.68（Assumed）。FY2024 0.863 含特別利益，不採用；第一季 0.788 含匯兌與營業外收益")

sec("G. 估值與對帳輸入")
for s in SC:
    inp(f"pk_{s}", f"情境 {s}：設備市場見頂年度（日曆年；循環位置規則用）", "年", None, PEAK[s], None, "Derived", PEAK_SRC[s], "設備股共用規則（tools/cycle_position.py）：FY(t) 早於見頂年＝中段、等於＝見頂、晚於＝谷底；設備市場模型的情境係數為 Assumed" if s == "A" else "同上")
inp("px", "股價（2026-10-23 收盤）", "¥", None, 15520, None, "Verified", "Yahoo!ファイナンス 時系列（2026-10-23 終值 15,520；前日 16,015）", "minkabu（2026-10-23 盤中 14,695，−8.24%）只確認當日走勢；收盤第二來源未找到")
inp("sa_fpe", "遠期 P/E（stockanalysis，2026-09-29）", "x", None, 21.86, None, "Interested-party", "stockanalysis.com statistics（股價 18,405）；EPS 年度未標明", "隱含遠期 EPS 約 842（Derived）")
inp("sa_ttm", "TTM P/E（stockanalysis，2026-09-29）", "x", None, 28.70, None, "Interested-party", "stockanalysis.com（TTM EPS 641.28）", "")
inp("cons_rev26", "共識營收 FY2026（2027年3月期）", "¥B", 177.8, 192.141, 196.7, "Interested-party", "minkabu アナリスト予想（2026-10-23 更新）：192,141 百萬日圓", "第二來源：stockanalysis（9 位，S&P Global）平均 189.7、區間 177.8–196.7（低／高取此）")
inp("cons_ni26", "共識淨利 FY2026", "¥B", None, 32.556, None, "Interested-party", "minkabu（2026-10-23）：32,556", "第二來源：未找到")
inp("cons_eps26", "共識 EPS FY2026", "¥", None, 801.46, None, "Interested-party", "minkabu（2026-10-23）：801.46", "第二來源：stockanalysis EPS 欄為「-」；kabuyoho 12 個月後前瞻 EPS 816.2（2026-07-30）")
inp("ks_12m", "kabuyoho：12 個月後前瞻 EPS（共識，2026-07-30）", "¥", None, 816.2, None, "Interested-party", "kabuyoho.jp reportTarget（資料日 2026-07-30）", "早於 8-04 指引上修")
inp("mtp_rev", "中期經營計畫：FY2027（2028年3月期）營收目標", "¥B", None, 185.0, None, "Interested-party", "公司 中期經營計畫（2025-05-09）：營收 1,850 億、營業利益 450 億、ROE 15%", "FY2026 指引營收已超過目標")
inp("mtp_oi", "中期經營計畫：FY2027 營業利益目標", "¥B", None, 45.0, None, "Interested-party", "同上", "")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}")
    if lo is not None: defname(key + "_lo", f"In!$D${r}")
    if hi is not None: defname(key + "_hi", f"In!$F${r}")
    r += 1
# 目標價參數（黃底；Assumed，Andy 可改）
r += 1
IN.cell(r, 1, "J. 目標價（DCF 與 Comparable 加權）與 FY2030–32 收斂法參數（黃底＝Assumed，Andy 可改）").font = BOLD; r += 1
def inrow(key, item, unit, base, tag, s1, s2, fmt=None, yellow=False):
    global r
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    c = IN.cell(r, 5, base)
    if not (isinstance(base, str) and base.startswith("=")): c.font = BLUE
    if fmt: c.number_format = fmt
    if yellow: c.fill = KEY
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}"); r += 1
for s_, w_ in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inrow(f"tp_w{s_}", f"目標價情境權重 {s_}（{SCN[s_]}）", "%", w_, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%", "0%", True)
inrow("tp_r", "折現率（股東權益要求報酬；美、台、日、韓同一值）", "%", 0.10, "Assumed", "Andy 可改（工作單預設 10%）", "日本利率較低，但依工作單不改值", "0.0%", True)
inrow("tp_g", "長期成長率（永續成長；收斂法 FY2032 營收年增的基準）", "%", 0.04, "Assumed", "Andy 可改（工作單預設 4%）", "情境 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt（PL 頁延伸區塊）", "0.0%", True)
inrow("tp_roe", "ROE（近三年平均，FY2023–FY2025）", "%", "=AVERAGE(0.128,0.155,0.135)", "Derived", "公司決算短信：FY2025 13.5%、FY2024 15.5%（" + TS25 + "）；FY2023 12.8%（roic.ai 12.79%）", "FY2023 公司說明資料擷取為 10.9%（擷取可能錯欄，未採用）", "0.0%")
inrow("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", 0.5, "Assumed", "Andy 可改（2026-10-25 11:53：DCF／Comparable 各半）", "—", "0%", True)
inrow("nim10", "十年平均淨利率（FY2016–FY2025，各年淨利 ÷ 營收的平均）", "%", "=AVERAGE(Hist!F7:F16)", "Derived", "Hist 頁（roic.ai 年度資料，FY2018–FY2025 與公司文件交叉）", "收斂法 B、C 情境的淨利率目標", "0.0%")
inrow("cons_tp", "市場共識目標價（賣方平均）", "¥", 20900, "Interested-party", "minkabu アナリスト予想（2026-10-23）：9 位平均 20,900（強力買進 3、買進 2、中立 3、強力賣出 1）", "第二來源：stockanalysis（S&P Global，9 位）平均 20,900、中位數 20,800、區間 16,500–24,000（資料更新 2026-06-24）", "#,##0")
inrow("cons_feps", "共識前瞻 EPS（差距拆解用）", "¥", "=cons_eps26", "Interested-party", "＝共識 FY2026（2027-03 期）EPS 801.46（minkabu 2026-10-23）；比本模型 12 個月後前瞻 EPS 早約一年半，EPS 因素因此偏大", "共識 FY2027 EPS：未找到（stockanalysis 需付費）", "#,##0")
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 80, 80]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"

# ---------------- Hist：FY2015–FY2025 年度實績 ----------------
H = wb.create_sheet("Hist")
H["A1"] = "Hist — 年度實績（十億日圓；EPS 為日圓）。合併：roic.ai 年度資料（來源1），FY2018 起與公司決算短信、說明資料交叉；部門（SPE、計測）：公司說明資料（FY2018 起）。用途：邊際利潤率迴歸（In 頁 F 節）、十年平均淨利率、十年 P/E 的 EPS"
H["A2"] = "會計年度 FY(t)＝t 年 4 月～t+1 年 3 月（例：FY2025＝2026年3月期）。過去十年無股票分割（各年 EPS 為公司公布值；roic.ai 加權股數 41–42 百萬股，穩定）"
hdr(H, 5, ["會計年度", "營收", "營業利益", "淨利", "EPS（基本，日圓）", "淨利率", "營業利益率", "標記", "SPE 營收", "SPE 營業利益", "計測營收", "計測營業利益", "來源1", "來源2"])
HIST = [("FY2015（2016-03）", 70.274, 13.219, 9.704, 234.58, None, "Derived（營業利益＝roic 營收 × 營益率 18.81%）", "roic.ai", "未找到"),
        ("FY2016（2017-03）", 77.792, 13.660, 9.909, 239.31, None, "Derived（營業利益＝營收 × 17.56%）", "roic.ai", "未找到"),
        ("FY2017（2018-03）", 88.194, 17.286, 12.717, 306.40, None, "Derived（營業利益＝營收 × 19.60%）", "roic.ai", "未找到"),
        ("FY2018（2019-03）", 101.520, 20.221, 14.665, 352.92, (69.117, 13.195, 32.403, 7.025), "Verified", "公司 FY2020/3 決算短信（英文，2020-05-11）", "roic.ai；部門：公司 FY2023/3 第三季說明資料（2023-02-06）"),
        ("FY2019（2020-03）", 87.927, 12.282, 7.156, 171.89, (56.198, 7.915, 31.728, 4.366), "Verified", "公司 FY2020/3 決算短信（2020-05-11）", "roic.ai"),
        ("FY2020（2021-03）", 97.105, 15.566, 12.175, 293.83, (71.745, 13.565, 25.359, 1.996), "Verified（營業利益 Derived：roic 營益率 16.03%；部門合計 15.561）", "roic.ai", "部門：公司 FY2023/3 第三季說明資料"),
        ("FY2021（2022-03）", 130.702, 28.327, 21.326, 522.52, (101.145, 24.698, 29.556, 3.628), "Verified", MATSUI, "公司 FY2023/3 決算說明資料（2023-05-11）；roic.ai"),
        ("FY2022（2023-03）", 146.801, 34.494, 23.630, 581.33, (112.365, 29.866, 34.436, 4.628), "Verified", MATSUI, "公司 FY2023/3 決算說明資料（2023-05-11）"),
        ("FY2023（2024-03）", 134.680, 25.307, 19.378, 480.49, (100.055, 19.899, 34.624, 5.408), "Verified", MATSUI, "公司 FY2024/3 決算說明資料（2024-05-10）"),
        ("FY2024（2025-03）", 150.534, 29.703, 25.637, 633.75, (113.481, 24.311, 37.053, 5.392), "Verified", MATSUI, FS25),
        ("FY2025（2026-03）", 166.839, 33.738, 24.739, 610.02, (127.878, 28.404, 38.960, 5.333), "Verified", TS25, MATSUI + "；" + FS25)]
for i, (fy, rv, oi, ni, eps, seg, tag, s1, s2) in enumerate(HIST, 6):
    H.cell(i, 1, fy)
    for j, v in enumerate((rv, oi, ni, eps), 2):
        H.cell(i, j, v).font = BLUE; H.cell(i, j).number_format = "#,##0.000" if j < 5 else "#,##0.00"
    H.cell(i, 6, f"=D{i}/B{i}").number_format = "0.0%"; H.cell(i, 7, f"=C{i}/B{i}").number_format = "0.0%"
    H.cell(i, 8, tag)
    if seg:
        for j, v in enumerate(seg, 9):
            H.cell(i, j, v).font = BLUE; H.cell(i, j).number_format = "#,##0.000"
    H.cell(i, 13, s1); H.cell(i, 14, s2)
H.cell(18, 1, "檢查：FY2025 列 − In 頁 a25_*、部門合計 − 合併（應約為 0）"); H["B18"] = "=ABS(B16-a25_rev)+ABS(C16-a25_oi)+ABS(D16-a25_ni)+ABS(E16-a25_eps)+ABS(I16+K16-B16)+ABS(J16+L16-C16)"; H["B18"].number_format = "0.000"
H.cell(20, 1, "迴歸與 Analogy（In 頁同值）").font = BOLD
for i, (lab, f, fm) in enumerate((("SPE 營業利益 ＝ 固定 ＋ 邊際 × 營收：邊際（FY2018–FY2025）", "=v_s", "0.000"), ("  R²", "=RSQ(J9:J16,I9:I16)", "0.000"),
                                  ("  固定（¥B）", "=INTERCEPT(J9:J16,I9:I16)", "0.00"),
                                  ("計測營業利益：邊際（FY2018–FY2025）", "=v_m", "0.000"), ("  R²", "=RSQ(L9:L16,K9:K16)", "0.000"),
                                  ("SPE 邊際（只用 FY2021–FY2025）", "=v_s2", "0.000"),
                                  ("計測營收年複合 FY2018→FY2025", "=(K16/K9)^(1/7)-1", "0.0%"), ("計測營收年複合 FY2020→FY2025", "=(K16/K11)^(1/5)-1", "0.0%"),
                                  ("SPE 營收年複合 FY2018→FY2025", "=(I16/I9)^(1/7)-1", "0.0%"), ("十年平均淨利率 FY2016–FY2025", "=nim10", "0.0%"),
                                  ("最大單年 SPE 營收下滑（FY2019）", "=I10/I9-1", "0.0%")), 21):
    H.cell(i, 1, lab); H.cell(i, 2, f).number_format = fm
H.column_dimensions["A"].width = 52
for c in "BCDEFGIJKL": H.column_dimensions[c].width = 12
H.column_dimensions["H"].width = 30; H.column_dimensions["M"].width = 50; H.column_dimensions["N"].width = 60

# ---------------- Q：季度實績與指引 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 季度實績（十億日圓；EPS 為日圓）。會計年度 4–3 月。來源1：松井證券季度表（百萬日圓）；來源2：公司決算說明資料（億日圓，部門）"
hdr(Q, 3, ["項目", "FY25 Q1", "FY25 Q2", "FY25 Q3", "FY25 Q4", "FY2025", "FY26 Q1", "FY26 Q2（公司上半指引隱含）", "標記", "來源1", "來源2／說明"])
QD = [("營收", [30.876, 46.194, 35.901, 53.868, None, 36.716, "=gh1_rev-q1_rev"], "#,##0.0", "松井：30,876／46,194／35,901／53,868；36,716", "公司上半指引 88.5 − 第一季"),
      ("營業利益", [4.581, 10.136, 6.215, 12.806, None, 5.920, "=19.1-q1_oi"], "#,##0.0", "松井：4,581／10,136／6,215／12,806；5,920", "上半指引 19.1"),
      ("淨利", [3.229, 6.383, 4.536, 10.591, None, 4.667, "=13.4-q1_ni"], "#,##0.0", "松井：3,229／6,383／4,536／10,591；4,667", "上半指引 13.4"),
      ("EPS（基本）", [79.77, 157.50, 111.88, 261.16, None, 115.01, "=330.14-115.01"], "#,##0.00", "松井", "上半指引 EPS 330.14"),
      ("SPE 營收", [23.5, 35.9, 27.4, 41.1, None, 28.5, None], "#,##0.0", BRQ1 + "（億日圓）", "FY2025 合計 127.9（說明資料四季合計 1,279）"),
      ("SPE 營業利益", [4.0, 8.3, 5.3, 10.7, None, 5.4, None], "#,##0.0", BRQ1, ""),
      ("SPE 受注", [26.4, 34.9, 26.0, 36.1, None, 53.2, None], "#,##0.0", BRQ1, "第一季 +102%：會社想定比多 100 億以上，約一半來自中國；第二季預期低於第一季（部分提前）"),
      ("計測營收", [7.3, 10.3, 8.5, 12.8, None, 8.2, None], "#,##0.0", BRQ1, ""),
      ("計測營業利益", [0.5, 1.8, 0.9, 2.1, None, 0.5, None], "#,##0.0", BRQ1, "")]
for i, (lab, vals, fm, s1, s2) in enumerate(QD, 4):
    Q.cell(i, 1, lab)
    for j, v in enumerate(vals):
        col = "BCDEFGH"[j]
        if v is None: continue
        c = Q[f"{col}{i}"]; c.value = v; c.number_format = fm
        if not isinstance(v, str): c.font = BLUE
    Q[f"F{i}"] = f"=SUM(B{i}:E{i})"; Q[f"F{i}"].number_format = fm
    Q.cell(i, 9, "Verified／Interested-party（指引）" if i <= 7 else "Interested-party（公司說明資料，億日圓取整）"); Q.cell(i, 10, s1); Q.cell(i, 11, s2)
Q.cell(13, 1, "營業利益率"); [Q.cell(13, j, f"={c}5/{c}4").__setattr__("number_format", "0.0%") for j, c in zip(range(2, 9), "BCDEFGH")]
Q.cell(14, 1, "檢查：FY2025 四季合計 − 年度實績（營收、營業利益、淨利；應約為 0）")
Q["B14"] = "=F4-a25_rev"; Q["C14"] = "=F5-a25_oi"; Q["D14"] = "=F6-a25_ni"
for c in "BCD": Q[f"{c}14"].number_format = "0.000"
Q.cell(15, 1, "第一季營收占公司全年指引"); Q["B15"] = "=q1_rev/g26_rev"; Q["B15"].number_format = "0.0%"; Q["C15"] = "FY2025 第一季占全年："; Q["D15"] = "=B4/a25_rev"; Q["D15"].number_format = "0.0%"
Q.cell(17, 1, "公司揭露的結構（" + QAQ1 + "；Interested-party）：第一季 HPC 占 SPE 營收 4 割半ば（邏輯 2／3、HBM 1／3）；HPC 受注約 200 億日圓台半（邏輯 3／4、HBM 1／4）；受注約為產能 1.5 倍；探針台、切割機交期 4–8 個月、研削機 6 個月以上")
Q.column_dimensions["A"].width = 40
for c in "BCDEFGH": Q.column_dimensions[c].width = 12
Q.column_dimensions["I"].width = 22; Q.column_dimensions["J"].width = 60; Q.column_dimensions["K"].width = 70

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
CYC = {y: chr(ord("C") + y - 2025) for y in range(2025, 2031)}   # C＝CY2025 … H＝CY2030
FYC = {y: chr(ord("C") + y - 2025) for y in range(2025, 2030)}   # C＝FY2025 … G＝FY2029
DEFP = dict(tail="g_tail", wpr="w_pr", wcmp="w_cmp", chn="chn", vs="v_s", vm="v_m", nioi="nioi", anc="g_anc", mspe="m_spe26", scn="A",
            exc=None, chh=None, gmet=None, madj=None)
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    exc = P["exc"] or f"exc_{s}"; chh = P["chh"] or f"chh_{s}"; gmet = P["gmet"] or f"gmet_{s}"
    madj = P["madj"] or ("madjC" if s == "C" else "0")
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "CY2025／FY2025", "CY2026／FY2026", "CY2027／FY2027", "CY2028／FY2028", "CY2029／FY2029", "CY2030", "", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 10, note); R[key] = r; r += 1
    ws.cell(r, 1, f"A. 設備市場（日曆年；markets/equipment 情境 {s}）").font = BOLD; r += 1
    for key, lab in (("tst", "測試與其他封裝設備"), ("apk", "先進封裝設備〔AI〕"), ("wfe", "晶圓前段設備")):
        row(key, lab, "US$B", {**{CYC[y]: f"={key}{s}_{y % 100}" for y in range(2025, 2030)}, "H": f"=G{r}*(1+{P['tail']})"}, fmt="#,##0.0",
            note="CY2030＝CY2029 ×（1＋g_tail）" if key == "tst" else "")
    ws.cell(r, 1, "B. 會計年度指數（FY(t)＝3/4 × CY(t)＋1/4 × CY(t+1)；FY2026＝1）").font = BOLD; r += 1
    fw = lambda key, y: f"(0.75*{CYC[y]}{R[key]}+0.25*{CYC[y+1]}{R[key]})"
    row("ti", "探針台需求指數（測試與其他封裝）", "x", {FYC[y]: f"={fw('tst', y)}/{fw('tst', 2026)}" for y in range(2025, 2030)}, fmt="0.000")
    row("mi", "加工裝置需求指數（前段 × w_cmp ＋ 先進封裝 ×（1 − w_cmp））", "x", {FYC[y]: f"={P['wcmp']}*{fw('wfe', y)}/{fw('wfe', 2026)}+(1-{P['wcmp']})*{fw('apk', y)}/{fw('apk', 2026)}" for y in range(2025, 2030)}, fmt="0.000")
    row("exf", "公司特有：超額成長係數（FY2027 起）", "x", {FYC[y]: f"=(1+{exc})^{y - 2026}" for y in range(2026, 2030)}, fmt="0.000")
    row("chf", "公司特有：中國係數（1 ＋ 中國占比 × 中國變化 × 生效比例）", "x", {"D": "=1", "E": f"=1+{P['chn']}*{chh}*0.5", "F": f"=1+{P['chn']}*{chh}", "G": f"=1+{P['chn']}*{chh}"}, fmt="0.000")
    ws.cell(r, 1, "C. 營收（FY2025 實績；FY2026 公司指引 ×（1＋g_anc）；FY2027 起驅動推導）").font = BOLD; r += 1
    spe26 = f"g26_spe*(1+{P['anc']})"
    row("pr", "探針台", "¥B", {"C": f"=spe_25*{P['wpr']}", "D": f"={spe26}*{P['wpr']}", **{FYC[y]: f"=$D{r}*{FYC[y]}{R['ti']}*{FYC[y]}{R['exf']}*{FYC[y]}{R['chf']}" for y in range(2027, 2030)}}, fmt="#,##0.0",
        note="FY2025 依 FY2026 上半構成比拆分（Assumed）")
    row("mc", "加工裝置（切割、研削、CMP 等）", "¥B", {"C": f"=spe_25*(1-{P['wpr']})", "D": f"={spe26}*(1-{P['wpr']})", **{FYC[y]: f"=$D{r}*{FYC[y]}{R['mi']}*{FYC[y]}{R['exf']}*{FYC[y]}{R['chf']}" for y in range(2027, 2030)}}, fmt="#,##0.0")
    row("spe", "半導體製造裝置（SPE）營收", "¥B", {FYC[y]: f"={FYC[y]}{R['pr']}+{FYC[y]}{R['mc']}" for y in range(2025, 2030)}, fmt="#,##0.0")
    row("met", "計測機器營收", "¥B", {"C": "=met_25", "D": f"=g26_met*(1+{P['anc']})", **{FYC[y]: f"={FYC[y-1]}{r}*(1+{gmet})" for y in range(2027, 2030)}}, fmt="#,##0.0")
    row("rev", "營收", "¥B", {FYC[y]: f"={FYC[y]}{R['spe']}+{FYC[y]}{R['met']}" for y in range(2025, 2030)}, fmt="#,##0.0", note="FY2026＝公司指引（g_anc＝0 時）")
    row("spsh", "SPE 占營收", "%", {FYC[y]: f"={FYC[y]}{R['spe']}/{FYC[y]}{R['rev']}" for y in range(2025, 2030)}, fmt="0.0%")
    ws.cell(r, 1, "D. 損益（部門營業利益＝FY2026 錨點 ＋ 邊際利潤率 × 營收變化）").font = BOLD; r += 1
    row("spop", "SPE 營業利益", "¥B", {"C": "=spe_op25", "D": f"=g26_spe*{P['mspe']}+{P['vs']}*(D{R['spe']}-g26_spe)",
                                     **{FYC[y]: f"=$D{r}+{P['vs']}*({FYC[y]}{R['spe']}-$D{R['spe']})" + (f"+{madj}*{FYC[y]}{R['spe']}" if y >= 2028 else "") for y in range(2027, 2030)}}, fmt="#,##0.0",
        note="情境 C 自 FY2028 另加 madjC × 營收")
    row("mtop", "計測營業利益", "¥B", {"C": "=met_op25", "D": f"=(g26_oi-g26_spe*{P['mspe']})+{P['vm']}*(D{R['met']}-g26_met)",
                                    **{FYC[y]: f"=$D{r}+{P['vm']}*({FYC[y]}{R['met']}-$D{R['met']})" for y in range(2027, 2030)}}, fmt="#,##0.0", note="FY2026＝指引營業利益 − SPE")
    row("oi", "營業利益", "¥B", {FYC[y]: f"={FYC[y]}{R['spop']}+{FYC[y]}{R['mtop']}" for y in range(2025, 2030)}, fmt="#,##0.0")
    row("spm", "SPE 營業利益率", "%", {FYC[y]: f"={FYC[y]}{R['spop']}/{FYC[y]}{R['spe']}" for y in range(2025, 2030)}, fmt="0.0%")
    row("oim", "營業利益率（結果）", "%", {FYC[y]: f"={FYC[y]}{R['oi']}/{FYC[y]}{R['rev']}" for y in range(2025, 2030)}, fmt="0.0%")
    row("ni", "淨利", "¥B", {"C": "=a25_ni", "D": f"=D{R['oi']}*g26_ni/g26_oi", **{FYC[y]: f"={FYC[y]}{R['oi']}*{P['nioi']}" for y in range(2027, 2030)}}, fmt="#,##0.0",
        note="FY2026＝營業利益 × 指引淨利÷營業利益；FY2027 起 × nioi")
    row("eps", "EPS", "¥", {"C": "=a25_eps", **{FYC[y]: f"={FYC[y]}{R['ni']}*1000/shares" for y in range(2026, 2030)}}, fmt="#,##0.0")
    row("chk", "檢查：FY2025 營收、營業利益 − 實績（部門合計捨入差約 0.002）", "¥B", {"C": f"=ABS(C{R['rev']}-a25_rev)+ABS(C{R['oi']}-a25_oi)"}, fmt="0.000")
    R["end"] = r
    return R

# ---------------- PL：三情境 ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 東京精密 損益（¥B；FY2025 實績；FY2026 錨定公司指引（2026-08-04 修正）；FY2027–FY2029 驅動推導；FY2030 起收斂法）。頂端為摘要，下方為三情境各一個完整推導區塊"
PB = {}; r0 = 60
TITLES = {"A": "情境 A（順利）：設備市場 A（見頂 2029）；SPE 超額成長 +3%/年；中國持平；計測 +6%",
          "B": "情境 B（平平）：設備市場 B（2028 持平、2029 −5%）；超額成長 0；中國 −5%；計測 +4%",
          "C": "情境 C（不如預期）：設備市場 C（2028 −20%）；超額成長 −3%/年；中國 −20%；計測 +1%；SPE 營益率 −1pt"}
for s in SC:
    PB[s] = engine(PL, r0, TITLES[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}（{SCN[s]}）").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "¥B", "#,##0.0"), ("pr", "  探針台", "¥B", "#,##0.0"), ("mc", "  加工裝置", "¥B", "#,##0.0"), ("met", "  計測機器", "¥B", "#,##0.0"),
                                ("oi", "營業利益", "¥B", "#,##0.0"), ("oim", "營業利益率", "%", "0.0%"), ("spm", "SPE 營業利益率", "%", "0.0%"), ("ni", "淨利", "¥B", "#,##0.0"), ("eps", "EPS", "¥", "#,##0")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in range(2025, 2030):
            PL[f"{FYC[y]}{sr}"] = f"={FYC[y]}{PB[s][key]}"; PL[f"{FYC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對：三情境 FY2025 − 實績（營收與營業利益；捨入差約 0.002）"); PL[f"C{sr}"] = "=" + "+".join(f"C{PB[s]['chk']}" for s in SC); PL[f"C{sr}"].number_format = "0.000"; sr += 1
PL.cell(sr, 1, "核對：FY2026 營收、營業利益、EPS − 公司指引（g_anc＝0 時應為 0；EPS 差為股數捨入）")
PL[f"C{sr}"] = f"=ABS(D{SUMR[('A', 'rev')]}-g26_rev)+ABS(D{SUMR[('A', 'oi')]}-g26_oi)"; PL[f"D{sr}"] = f"=D{SUMR[('A', 'eps')]}-g26_eps"
PL[f"C{sr}"].number_format = "0.000"; PL[f"D{sr}"].number_format = "0.00"; CHK26 = sr
# 延伸到 FY2032（收斂法）
p0 = PB["C"]["end"] + 3
PL.cell(p0, 1, "延伸到 FY2032：FY2025–FY2029 引用本頁頂端摘要；灰底 FY2030–FY2032 為收斂法，非驅動因子推導").font = BOLD
PL.cell(p0 + 1, 1, "收斂法：營收年增由 FY2029 年增線性收斂到 FY2032 的長期成長（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt）；淨利率由 FY2029 線性收斂（A＝FY2029、B＝FY2029 與十年平均的中點、C＝十年平均 nim10）；股數維持（EPS 與淨利同比例）")
EC = {y: chr(ord("C") + y - 2025) for y in range(2025, 2033)}
hdr(PL, p0 + 2, ["情境（延伸）", "單位", "FY2025A", "FY2026", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
EXT = {}; pr = p0 + 3; gadd = {"A": "+0.02", "B": "", "C": "-0.02"}
for s in SC:
    PL.cell(pr, 1, f"情境 {s}（{SCN[s]}）").font = BOLD; pr += 1
    R = {k: pr + i for i, k in enumerate(("rev", "g", "nim", "ni", "eps"))}
    for k, lab in (("rev", "營收"), ("g", "  營收年增"), ("nim", "  淨利率（淨利 ÷ 營收）"), ("ni", "淨利（收斂法延伸）"), ("eps", "EPS（延伸）")):
        PL.cell(R[k], 1, lab); PL.cell(R[k], 2, {"rev": "¥B", "ni": "¥B", "eps": "¥"}.get(k, "%"))
    for y in range(2025, 2033):
        c = EC[y]; pc = EC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={FYC[y]}{SUMR[(s, 'rev')]}"; PL[f"{c}{R['ni']}"] = f"={FYC[y]}{SUMR[(s, 'ni')]}"
            PL[f"{c}{R['eps']}"] = f"={FYC[y]}{SUMR[(s, 'eps')]}"; PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
        else:
            k = y - 2029
            mT = {"A": f"$G{R['nim']}", "B": f"($G{R['nim']}+nim10)/2", "C": "nim10"}[s]
            PL[f"{c}{R['g']}"] = f"=$G{R['g']}+((tp_g{gadd[s]})-$G{R['g']})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({mT}-$G{R['nim']})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['eps']}"] = f"=$G{R['eps']}*{c}{R['ni']}/$G{R['ni']}"
            for kk in R: PL[f"{c}{R[kk]}"].fill = FADE
        for kk, fm in (("rev", "#,##0.0"), ("g", "0.0%"), ("nim", "0.0%"), ("ni", "#,##0.0"), ("eps", "#,##0")):
            if PL[f"{c}{R[kk]}"].value is not None: PL[f"{c}{R[kk]}"].number_format = fm
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導")
    EXT[s] = R; pr += 6
PL.column_dimensions["A"].width = 56; PL.column_dimensions["B"].width = 10; PL.column_dimensions["J"].width = 46; PL.column_dimensions["K"].width = 36
for c in "CDEFGHI": PL.column_dimensions[c].width = 13

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — 估值（日圓）。目標價＝DCF 與 Comparable 加權（第 0 節）；跨公司規則值（十年年底 P/E 依循環位置取分位：見頂 25 分位、中段中位數、谷底 75 分位）保留為參考（第 5–6 節）"
V["A2"] = "P/E 定義（自算）：日曆年底收盤價 ÷ 「當期會計年度」公司公布的基本 EPS（例：2025-12 對 FY2025＝2026年3月期）。過去十年無股票分割"
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(rr, key, label, fs, fmt, note=""):
    V.cell(rr, 1, label); T[key] = rr
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{rr}"] = f; V[f"{c}{rr}"].number_format = fmt
    if note: V[f"E{rr}"] = note
t0 = 4
V.cell(t0, 1, "0. 目標價").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
V.cell(t0 + 3, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{t0 + 3}"] = "PL 頁延伸區塊；FY2030 起為收斂法"
EPSR = {}
for i, y in enumerate(YRS6):
    tset(t0 + 4 + i, f"eps{y}", f"  FY{y}", [f"=PL!{EC[y]}{EXT[s]['eps']}" for s in SC], "#,##0"); EPSR[y] = t0 + 4 + i
rr = t0 + 10
V.cell(rr, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); V[f"E{rr}"] = "模型無折舊與資本支出項目，故不用（淨利＋折舊－資本支出）；公司 FY2026 資本支出指引 150 億、折舊 62 億（擴產期），近年股東現金流可能低於此式（偏高估）"
tset(rr + 1, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
RV = rr + 1; CFR = {}
for i, y in enumerate(YRS6):
    tset(rr + 2 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], "#,##0"); CFR[y] = rr + 2 + i
c1, c6 = CFR[2027], CFR[2032]
tset(rr + 8, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], "#,##0",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027；低於 0 取 0；會計年度 4–3 月，FY2027＝2027-04～2028-03 仍以 t＝1 折現（略高估）")
tset(rr + 9, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{rr + 8}*(1+tp_r)" for c in "BCD"], "#,##0")
tset(rr + 10, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{rr + 9}+tp_wB*C{rr + 9}+tp_wC*D{rr + 9}", None, None], "#,##0", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁 tp_wA／tp_wB／tp_wC")
q = rr + 12
V.cell(q, 1, "0.2 Comparable（市場主流同業前瞻 P/E；非 A 股公司，同業不放 A 股）").font = BOLD
hdr(V, q + 1, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("Disco", "6146.T", 26.01, "2026-10-23", "stockanalysis.com statistics：Forward PE（分析師共識；價格日期如左）"),
         ("東京威力科創（Tokyo Electron）", "8035.T", 27.09, "2026-10-21", "同上"),
         ("Advantest", "6857.T", 37.69, "2026-10-21", "同上"),
         ("Lasertec", "6920.T", 34.02, "2026-10-23", "同上"),
         ("Kokusai Electric", "6525.T", 35.83, "2026-10-02", "同上（價格日期較舊）")]
for i, (nm, tk, pe, dt, src) in enumerate(PEERS, q + 2):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
p1 = q + 2; pl_ = q + 1 + len(PEERS)
V.cell(pl_ + 1, 1, "並列參考（不進目標價）：東京精密自身前瞻 P/E"); V.cell(pl_ + 1, 3, "=sa_fpe").number_format = "0.0"; V.cell(pl_ + 1, 4, "2026-09-29"); V.cell(pl_ + 1, 5, "stockanalysis.com；本公司長期低於同業（見第 2 節十年 P/E）")
q = pl_ + 2
tset(q, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{p1}:C{pl_})", None, None], "0.0", "日本半導體設備同業：Disco（研削／切割同業）、東京威力科創（探針台同業）、Advantest（測試）、Lasertec、Kokusai Electric")
tset(q + 1, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*(5/12*{c}{EPSR[2027]}+7/12*{c}{EPSR[2028]})" for s, c in zip(SC, "BCD")), None, None], "#,##0",
     "約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 5／12 × FY2027（2027-04～2028-03）＋7／12 × FY2028（會計年度 4–3 月，依月份調整）")
tset(q + 2, "comptp", "Comparable 目標價", [f"=B{q}*B{q + 1}", None, None], "#,##0", "同業中位數 × 上列 EPS")
q += 4
V.cell(q, 1, "0.3 結論").font = BOLD; T["h03"] = q
hdr(V, q + 1, ["項目", "數值", "", "", "說明"]); T["hd3"] = q + 1
q += 2
tset(q, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], "#,##0", "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁")
tset(q + 1, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], "#,##0", "引用 0.2")
tset(q + 2, "cons", "市場共識目標價", ["=cons_tp", None, None], "#,##0", "minkabu（2026-10-23）：9 位分析師平均 20,900；stockanalysis（S&P Global）同值")
tset(q + 3, "upw", "上檔空間：加權目標價", [f"=B{q}/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-23 收盤 15,520）")
tset(q + 4, "upc", "上檔空間：Comparable 目標價", [f"=B{q + 1}/px-1", None, None], "0.0%")
tset(q + 5, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型 12 個月後前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝FY2026 共識 801.46（Recon），早約一年半，EPS 因素因此偏大")
tset(q + 6, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%")
V[f"E{q + 6}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(q + 7, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
q8 = q + 8
T["pr"] = q8  # 公式在 P/E 表建好後填入
V.cell(q8, 1, "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）")
V.cell(q + 9, 1, "目標價方法"); T["meth"] = q + 9
V[f"B{q + 9}"] = ("DCF：股東現金流＝EPS ×（1 − g／ROE）（模型無折舊與資本支出項目），r 10%、g 4%、ROE 為近三年平均，FY2027–32 明確期（FY2030 起收斂法）＋永續。"
                  "Comparable：Disco、東京威力科創、Advantest、Lasertec、Kokusai 前瞻 P/E 中位數 × 12 個月後前瞻 EPS。"
                  "偏誤：本公司十年 P/E 與自身前瞻 P/E 長期低於同業（探針台、計測的成長與利潤率低於同業），Comparable 偏高估；DCF 未扣擴產期資本支出，偏高估。")
for k in ("dcftp", "comptp", "tpw"): V[f"B{T[k]}"].fill = KEY
# 1. 價格與市場隱含倍數
b0 = q + 11
V.cell(b0, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, b0 + 1, ["項目", "A", "B", "C", "說明"])
fr = b0 + 2
V.cell(fr, 1, "現價（2026-10-23 收盤）"); V.cell(fr, 2, "=px").number_format = "#,##0"; V.cell(fr, 5, "Yahoo!ファイナンス"); T["px"] = fr; fr += 1
V.cell(fr, 1, "現價 ÷ FY2026 EPS（錨點）")
for j, s in enumerate(SC): V.cell(fr, 2 + j, f"=px/PL!D{SUMR[(s, 'eps')]}").number_format = "0.0"
V.cell(fr, 5, "FY2026 三情境相同（公司指引）"); fr += 1
for y in (2027, 2028, 2029):
    V.cell(fr, 1, f"現價 ÷ FY{y} EPS（模型）")
    for j, s in enumerate(SC): V.cell(fr, 2 + j, f"=px/PL!{FYC[y]}{SUMR[(s, 'eps')]}").number_format = "0.0"
    T[f"pe{y}"] = fr; fr += 1
for nm, f, note in (("P/E：現價 ÷ FY2025 EPS 610.02", "=px/a25_eps", "上一個會計年度"),
                    ("P/E：現價 ÷ 近 12 月 EPS（FY25 Q2–Q4＋FY26 Q1）", "=px/(157.5+111.88+261.16+115.01)", "近 12 月 EPS ＝ 157.50＋111.88＋261.16＋115.01 ＝ 645.55（Derived，Q 頁）"),
                    ("TTM P/E：stockanalysis（2026-09-29）", "=sa_ttm", "參考"), ("遠期 P/E：stockanalysis（共識，2026-09-29）", "=sa_fpe", "參考，不用於估值")):
    V.cell(fr, 1, nm); V.cell(fr, 2, f).number_format = "0.0"; V.cell(fr, 5, note); fr += 1
# 2. 十年 P/E
p0v = fr + 1
V.cell(p0v, 1, "2. 10 年 P/E（年底）").font = BOLD
hdr(V, p0v + 1, ["年底", "收盤價（未調整）", "來源2：Yahoo Finance 收盤", "兩者差距", "", "", "當期 FY EPS", "P/E 當期", "標記", "說明"])
PX = {2016: (3465, 3465, "2016-12-30"), 2017: (4460, 4460, "2017-12-29"), 2018: (2773, 2773, "2018-12-28"), 2019: (4260, 4260, "2019-12-30"),
      2020: (4850, 4850, "2020-12-30"), 2021: (5100, 5100, "2021-12-30"), 2022: (4270, 4270, "2022-12-30"), 2023: (8671, 8671, "2023-12-29"),
      2024: (7357, 7357, "2024-12-30"), 2025: (11105, 11105, "2025-12-30")}
y0 = p0v + 2
for i, y in enumerate(range(2016, 2026)):
    rr_ = y0 + i; raw, yus, d = PX[y]
    V.cell(rr_, 1, f"{y}-12（{d}）"); V.cell(rr_, 2, raw).font = BLUE; V.cell(rr_, 3, yus).font = BLUE
    V.cell(rr_, 4, f"=B{rr_}/C{rr_}-1").number_format = "0.00%"; V.cell(rr_, 7, f"=Hist!E{7 + i}").number_format = "#,##0.00"
    V.cell(rr_, 8, f"=B{rr_}/G{rr_}").number_format = "0.0"; V.cell(rr_, 2).number_format = "#,##0"; V.cell(rr_, 3).number_format = "#,##0"
    V.cell(rr_, 9, "Verified"); V.cell(rr_, 10, "來源1：Yahoo!ファイナンス 時系列（終值）；來源2：Yahoo Finance 月線收盤；EPS：Hist 頁" if i == 0 else ("2019 年 P/E 偏高：FY2019 EPS 因設備下行減半（171.89）" if y == 2019 else ""))
y9 = y0 + 9; PE = f"H{y0}:H{y9}"
V.cell(y9 + 1, 1, "兩個價格來源可能共用同一資料商（不完全獨立）；10 年都有價格，無缺年。調整後收盤（含股利再投資）不用於 P/E。")
s0v = y9 + 3
V.cell(s0v, 1, "3. 分位（10 年）").font = BOLD
hdr(V, s0v + 1, ["統計", "P/E 當期", "用途"])
for i, (nm, p, use) in enumerate((("25 分位", 0.25, "見頂年"), ("中位數", 0.5, "中段"), ("75 分位", 0.75, "谷底")), s0v + 2):
    V.cell(i, 1, nm); V.cell(i, 2, f"=PERCENTILE({PE},{p})").number_format = "0.0"; V.cell(i, 3, use)
defname("pe_peak", f"Val!$B${s0v + 2}"); defname("pe_mid", f"Val!$B${s0v + 3}"); defname("pe_trough", f"Val!$B${s0v + 4}")
V.cell(s0v + 5, 1, "年數"); V[f"B{s0v + 5}"] = f"=COUNT({PE})"
T["pct"] = s0v
# 0.3 的分位檢查
V[f"B{q8}"], V[f"C{q8}"], V[f"D{q8}"] = [f"=IF(px/{c}{EPSR[2028]}<MIN({PE}),0,IF(px/{c}{EPSR[2028]}>MAX({PE}),1,PERCENTRANK({PE},px/{c}{EPSR[2028]})))" for c in "BCD"]
for c in "BCD": V[f"{c}{q8}"].number_format = "0%"
V[f"E{q8}"] = "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 2 節當期口徑）的分位；只作檢查，不進目標價"
# 4. 循環位置
c0 = s0v + 7
V.cell(c0, 1, "4. FY2027 的循環位置（設備股共用規則；見頂年度讀設備市場模型，不在本模型另訂）").font = BOLD
V.cell(c0 + 1, 1, CP.rule_text(PEAK))
hdr(V, c0 + 2, ["情境", "設備市場見頂年度", "FY2027 對照日曆年", "判定", "倍數"])
CYR = {}
for i, s in enumerate(SC, c0 + 3):
    V.cell(i, 1, s); V.cell(i, 2, f"=pk_{s}"); V.cell(i, 3, 2027).font = BLUE
    V.cell(i, 4, "=" + CP.excel_rule(f"C{i}", f"pk_{s}"))
    V.cell(i, 5, f'=IF(D{i}="見頂",pe_peak,IF(D{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"; CYR[s] = i
T["cyc"] = c0
# 5. 規則值隱含股價
v0 = c0 + 7
V.cell(v0, 1, "5. 規則值隱含股價（FY2027 EPS × 依判定的倍數；其餘倍數僅供參考）").font = BOLD
hdr(V, v0 + 1, ["情境", "EPS 年度", "EPS", "倍數類型", "倍數", "隱含股價", "相對現價", "備註"])
vr = v0 + 2; VAL = {}
for i, s in enumerate(SC):
    for y, f, nm in ((2027, f"E{CYR[s]}", None), (2027, "pe_peak", "參考：25 分位"), (2027, "pe_mid", "參考：中位數"), (2027, "pe_trough", "參考：75 分位"), (2028, "pe_mid", "參考：FY2028 × 中位數")):
        V.cell(vr, 1, s); V.cell(vr, 2, f"FY{y}"); V.cell(vr, 3, f"=PL!{FYC[y]}{SUMR[(s, 'eps')]}").number_format = "#,##0"
        V.cell(vr, 4, nm if nm else f"=\"規則值：\"&D{CYR[s]}"); V.cell(vr, 5, f"={f}").number_format = "0.0"
        V.cell(vr, 6, f"=C{vr}*E{vr}").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
        if nm is None:
            for cc in range(1, 9): V.cell(vr, cc).fill = KEY
            VAL[s] = vr
        vr += 1
T["v0"] = v0
# 6. 舊法：規則值情境加權（參考）
o0 = vr + 1
V.cell(o0, 1, "6. 規則值情境加權（參考，不是目標價）").font = BOLD
hdr(V, o0 + 1, ["項目", "數值", "", "", "說明"])
V.cell(o0 + 2, 1, "規則值情境加權價（12 個月；參考）"); V.cell(o0 + 2, 2, f"=tp_wA*F{VAL['A']}+tp_wB*F{VAL['B']}+tp_wC*F{VAL['C']}").number_format = "#,##0"
V.cell(o0 + 3, 1, "相對現價（規則值情境加權價 ÷ 現價 − 1）"); V.cell(o0 + 3, 2, f"=B{o0 + 2}/px-1").number_format = "0.0%"
V.cell(o0 + 4, 1, "規則值方法"); V.cell(o0 + 4, 2, "各情境 FY2027 EPS × 十年（2016–2025 年底）P/E 依循環位置取分位；十年倍數含 2018–2022 年設備股低估值期（7–17x），偏誤方向為低估")
T["old"] = o0
V.column_dimensions["A"].width = 52
for c in "BCDFGHI": V.column_dimensions[c].width = 14
V.column_dimensions["E"].width = 60; V.column_dimensions["J"].width = 60

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A）：對 FY2027、FY2028 營收與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方，與 PL 同一引擎）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收 ¥B", "FY2027 EPS ¥", "相對基準", "FY2028 營收 ¥B", "FY2028 EPS ¥", "相對基準", "FY2026 EPS ¥", "相對基準"])
cases = [("SPE 超額成長 +1%/年", "exc_A 低", dict(exc="exc_A_lo")), ("SPE 超額成長 +5%/年", "exc_A 高", dict(exc="exc_A_hi")),
         ("超額成長 0（與市場同步）", "exc＝0", dict(exc="0")),
         ("中國 SPE −20%（情境 C 的值）", "chh＝chh_C", dict(chh="chh_C")),
         ("中國占比 40%", "chn 低（搭配 chh_C）", dict(chn="chn_lo", chh="chh_C")), ("中國占比 50%", "chn 高（搭配 chh_C）", dict(chn="chn_hi", chh="chh_C")),
         ("探針台占 SPE 70%", "w_pr 低", dict(wpr="w_pr_lo")), ("探針台占 SPE 74%", "w_pr 高", dict(wpr="w_pr_hi")),
         ("加工裝置前段比例 10%", "w_cmp 低", dict(wcmp="w_cmp_lo")), ("加工裝置前段比例 35%", "w_cmp 高", dict(wcmp="w_cmp_hi")),
         ("CY2030 設備市場 −10%", "g_tail 低", dict(tail="g_tail_lo")), ("CY2030 設備市場 +10%", "g_tail 高", dict(tail="g_tail_hi")),
         ("計測年增 4%", "gmet_A 低", dict(gmet="gmet_A_lo")), ("計測年增 8%", "gmet_A 高", dict(gmet="gmet_A_hi")),
         ("SPE 邊際利潤率只用 FY2021–25", "v_s2", dict(vs="v_s2")),
         ("淨利÷營業利益 0.68", "nioi 低", dict(nioi="nioi_lo")), ("淨利÷營業利益 0.733（FY2025）", "nioi 高", dict(nioi="nioi_hi")),
         ("FY2026 實績低於指引 3%", "g_anc 低", dict(anc="g_anc_lo")), ("FY2026 實績高於指引 5%", "g_anc 高", dict(anc="g_anc_hi"))]
def put_case(i, name, setting, R, pre=""):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    SE[f"C{i}"] = f"={pre}{FYC[2027]}{R['rev']}"; SE[f"D{i}"] = f"={pre}{FYC[2027]}{R['eps']}"; SE[f"F{i}"] = f"={pre}{FYC[2028]}{R['rev']}"; SE[f"G{i}"] = f"={pre}{FYC[2028]}{R['eps']}"
    SE[f"E{i}"] = f"=D{i}/D$4-1"; SE[f"H{i}"] = f"=G{i}/G$4-1"
    SE[f"I{i}"] = f"={pre}{FYC[2026]}{R['eps']}"; SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c in "CDFGI": SE[f"{c}{i}"].number_format = "#,##0"
    for c in "EHJ": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
SB = {}; base_row = 60
for n, (nm, setting, ov) in enumerate(cases):
    SB[nm] = engine(SE, base_row + n * 40, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    put_case(5 + n, nm, setting, SB[nm])
nn = 5 + len(cases)
put_case(nn, "情境 B", "設備市場 B；超額 0；中國 −5%；計測 +4%", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "設備市場 C；超額 −3%；中國 −20%；計測 +1%；營益率 −1pt", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2026 錨定公司指引，只受 g_anc 影響；驅動參數從 FY2027 起作用。FY2026 SPE 營益率（m_spe26）只改變部門拆分，不改變合併營業利益與 EPS，故不列案例；g_tail 只影響 FY2029（CY2030 的 1/4）。")
# 目標價參數
s0 = base_row + len(cases) * 40 + 3
SE.cell(s0, 1, "目標價參數（只改 DCF 參數與權重，PL 收斂法的長期成長不隨 tp_g 變動）").font = BOLD
hdr(SE, s0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def dcf_formula(c, r_="tp_r", g="tp_g"):
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-MIN({g}/tp_roe,0.8))"
    pv = "+".join(f"{cf(y)}/(1+{r_})^{t}" for t, y in enumerate(YRS6, 1))
    return f"MAX(0,{pv}+{cf(2032)}*(1+{g})/({r_}-{g})/(1+{r_})^6)"
tcases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r_="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r_="(tp_r-0.01)")),
          ("長期成長 +1pt", "g 5%（永續與再投資率）", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
          ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(tcases):
    i = s0 + 2 + k; w = kw.pop("w", "tp_w_dcf"); r_ = kw.get("r_", "tp_r")
    SE.cell(i, 1, nm); SE.cell(i, 2, st)
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, **kw)}" for s, c in zip(SC, "BCD")) + f")*(1+{r_})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"; SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"; SE[f"F{i}"] = f"=E{i}/E${s0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "#,##0"
    SE[f"F{i}"].number_format = "0.0%"
TS0 = s0
SE.column_dimensions["A"].width = 40; SE.column_dimensions["B"].width = 40
for c in "CDEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — FY2026 錨定公司指引（A 節為恆等式，不是檢驗）；B 節共識只作參考，未用於校準或錨點；C 節為結構檢查。差距＝模型 ÷ 對照 − 1"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False):
    global rr_
    RC.cell(rr_, 1, label); RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        L_ = chr(64 + cv)
        RC.cell(rr_, cd, f"={L_}{rr_}-B{rr_}" if pt else f"={L_}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y=2026: (lambda s: f"=PL!{FYC[y]}{SUMR[(s, key)]}")
RC.cell(rr_, 1, "A. FY2026 錨點（公司指引 2026-08-04）：差距為 0 是恆等式").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 公司指引", "=g26_rev", pl("rev"), "決算短信（第一季）189.0")
rec("FY2026 營業利益 vs 公司指引", "=g26_oi", pl("oi"), "44.0")
rec("FY2026 EPS vs 公司指引", "=g26_eps", pl("eps"), "758.83", "差異為股數捨入", fmt="#,##0.00")
rec("FY2026 營業利益率 vs FY2025 實績（百分點）", "=a25_oi/a25_rev", pl("oim"), "FY2025 20.2%", "公司稱高溫控探針台、HRG 研削機、切割機量增改善組合", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "A2. 第一季進度（只作合理性檢查）").font = BOLD; rr_ += 1
rec("FY2026 上半指引隱含第二季營收 vs 第一季實績", "=q1_rev", lambda s: "=gh1_rev-q1_rev", "上半指引 88.5 − 第一季 36.7", "第二季須 +41%；FY2025 第二季對第一季 +50%（Q 頁）")
rec("FY2026 下半營收（指引隱含）vs FY2025 下半實績", "=a25_rev-77.07", lambda s: f"=PL!D{SUMR[(s, 'rev')]}-gh1_rev", "FY2025 上半 77.07（松井）", "下半 100.5 vs 89.8")
RC.cell(rr_, 1, "B. FY2026 vs 共識（參考）").font = BOLD; rr_ += 1
rec("營收 vs 共識（minkabu 2026-10-23）", "=cons_rev26", pl("rev"), "minkabu 192.1；stockanalysis 189.7（區間 177.8–196.7）")
rec("淨利 vs 共識（minkabu）", "=cons_ni26", pl("ni"), "minkabu 32.6")
rec("EPS vs 共識（minkabu）", "=cons_eps26", pl("eps"), "minkabu 801.46", "共識高於公司指引 5.6%：市場預期再上修", fmt="#,##0")
rec("FY2027 EPS vs kabuyoho 12 個月後前瞻 EPS（2026-07-30；口徑不同，參考）", "=ks_12m", pl("eps", 2027), "kabuyoho 816.2", "早於 8-04 上修；12 個月後≈ FY2026 與 FY2027 混合", fmt="#,##0")
rec("FY2027 營收 vs 中期經營計畫目標（185.0）", "=mtp_rev", pl("rev", 2027), "公司 中期經營計畫（2025-05-09）", "目標訂於 AI 需求加速前")
rec("FY2027 營業利益 vs 中期經營計畫目標（45.0）", "=mtp_oi", pl("oi", 2027), "同上")
RC.cell(rr_, 1, "C. 結構檢查").font = BOLD; rr_ += 1
rec("FY2026 SPE 營收年增 vs 測試與其他封裝設備 CY2026 年增（百分點）", f"=tstA_26/tstA_25-1", lambda s: f"=PL!D{PB[s]['spe']}/PL!C{PB[s]['spe']}-1", "markets/equipment（情境 A；2026 含口徑轉換）", "公司指引 SPE +15.3%，低於市場 +26.6%：探針台產能滿載（受注約為產能 1.5 倍）", fmt="0.0%", pt=True)
rec("FY2028 SPE 營收 ÷ FY2026（模型）vs 設備總額 CY2028 ÷ CY2026", f"=0", lambda s: f"=PL!F{PB[s]['spe']}/PL!D{PB[s]['spe']}-1-(tot_{s}_28/tot_{s}_26-1)", "markets/equipment 設備總額（同字母）", "差距欄＝模型 − 市場（百分點）。模型低於市場的原因：會計年度比日曆年晚一季、測試與封裝設備成長低於前段、超額成長與中國變數", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "參考：HPC 占 SPE 營收（FY2026 第一季，公司）"); RC.cell(rr_, 3, "=hpc").number_format = "0%"; RC.cell(rr_, 10, "邏輯 2／3、HBM 1／3；上游設備路徑已含 AI，未另行加計"); rr_ += 1
RC.column_dimensions["A"].width = 66
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 40; RC.column_dimensions["J"].width = 70

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文以網頁擷取工具讀取；表格錯位處以算術核對；版面與頁碼待人工複核）"; SO["A1"].font = BOLD
srcs = [("JPX：2026年3月期 決算短信（2026-05-13）", "https://www2.jpx.co.jp/disc/77290/081220260518539874_tse-acedjpsm-77290-20260518377290-ixbrl.htm"),
 ("JPX：2027年3月期 第1四半期 決算短信（2026-08-04；業績予想修正）", "https://www2.jpx.co.jp/disc/77290/081220260803507651_tse-qcedjpsm-77290-20260804377290-ixbrl.htm"),
 ("公司：2026年3月期 決算補足資料（FS_2026_4q_J_rev）", "https://www.accretech.com/jp/ir/news/news-202605180003/main/0/link/FS_2026_4q_J_rev.pdf"),
 ("公司：2026年度第1四半期 決算說明資料（2026-08-04）", "https://www.accretech.com/jp/ir/news/auto_20260804508092/main/0/link/BR_2027-3_1q_J_2.pdf"),
 ("公司：2026年度第1四半期 Q&A（2026-08-04）", "https://www.accretech.com/jp/ir/news/news-202608041630/main/0/link/QA_2027-3_1q_J.pdf"),
 ("公司：2025年度第1四半期 Q&A（2025-08-04）", "https://www.accretech.com/jp/ir/news/news3574518910588001242/main/0/link/QA_2026_1q_J.pdf"),
 ("公司：中期經營計畫（2025-05-09；部門 FY2022/3–FY2025/3）", "https://www.accretech.com/en/ir/news/auto_20250509537325/pdfFile.pdf"),
 ("公司：FY2024/3 決算說明資料（2024-05-10）", "https://www.accretech.com/en/ir/news/auto_20240509587976/pdfFile.pdf"),
 ("公司：FY2023/3 決算說明資料（2023-05-11）", "https://www.accretech.com/en/ir/news/auto_20230511567918/main/0/link/BR_2023_4q_E_r.pdf"),
 ("公司：FY2023/3 第三季說明資料（2023-02-06；部門 FY2019/3–FY2022/3）", "https://www.accretech.com/en/ir/news/auto_20230206501370/pdfFile.pdf"),
 ("公司：FY2020/3 決算短信（英文，2020-05-11）", "https://www.accretech.com/en/ir/news/auto_20200510407802/main/0/link/FS_2020_4q_E_2.pdf"),
 ("松井證券：決算情報（年度、季度）", "https://finance.matsui.co.jp/stock/7729/settlement/index"),
 ("roic.ai：年度財務（FY2014/3–FY2026/3）", "https://www.roic.ai/quote/7729.T/classic"),
 ("Yahoo!ファイナンス：株價時系列（年底收盤、2026-10-23）", "https://finance.yahoo.co.jp/quote/7729.T/history"),
 ("Yahoo Finance：Historical prices（月線）", "https://finance.yahoo.com/quote/7729.T/history/"),
 ("minkabu：アナリスト予想（2026-10-23）", "https://minkabu.jp/stock/7729/analyst_consensus"),
 ("stockanalysis：共識、P/E（7729、6146、8035、6857、6920、6525）", "https://stockanalysis.com/quote/tyo/7729/forecast/"),
 ("kabuyoho：目標株價、12 個月後前瞻 EPS（2026-07-30）", "https://kabuyoho.jp/reportTarget?bcode=7729"),
 ("市場層：markets/equipment 成品（本 repo）；tools/cycle_position.py", ""),
 ("未找到：探針台、切割機市占一手調查；切割／研削／CMP 個別金額；中國營收占比；共識 FY2027 EPS；FY2015–FY2017 部門資料", ""),
 ("無法開啟（建置環境被擋或逾時）：finboard.jp、gyokaidigest、kabuyoho 共識頁、monex scouter（速率限制）", ""),
 ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 80; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
PROP = ("東京精密約八成營收來自半導體設備（探針台約七成、切割／研削／CMP 約三成），其餘為計測機器；HPC 已占設備營收四成半。"
        "問題是 AI 測試與先進封裝的增量能否撐過設備循環，並讓現價約 20 倍的公司指引本益比（低於日本設備同業的 30 倍以上）被重估？")
lines = [("東京精密（7729.T）收支與估值模型 v0.1.1（AI 半導體 Project：設備與材料）　2026-10-25", True),
 ("模型命題：" + PROP, False),
 ("（約 20 倍＝現價 ¥15,520（2026-10-23 收盤）÷ 公司 FY2026 指引 EPS ¥758.83 ＝ 20.5 倍；見 Val 頁第 1 節）", False), ("", False),
 ("驅動因子 → 推導量（因果方向：左邊決定右邊）", True),
 ("PL 區塊｜markets/equipment 同字母情境的「測試與其他封裝設備」（CY）→ FY(t)＝3/4 × CY(t)＋1/4 × CY(t+1) → 探針台需求指數（FY2026＝1）× FY2026 探針台營收 → 探針台營收", False),
 ("PL 區塊｜「晶圓前段設備」× w_cmp ＋「先進封裝設備〔AI〕」×（1 − w_cmp）→ 加工裝置需求指數 × FY2026 加工裝置營收 → 加工裝置營收", False),
 ("PL 區塊｜公司特有：超額成長（A +3%／B 0／C −3% 每年）、中國變化（A 0／B −5%／C −20%，× 中國占比 45%）→ SPE 營收", False),
 ("PL 區塊｜計測機器：FY2026 指引 41.5 ×（1＋年增；A 6%／B 4%／C 1%，Analogy：公司歷史）", False),
 ("Hist → PL｜部門營業利益＝FY2026 錨點 ＋ 邊際利潤率（FY2018–FY2025 迴歸）× 營收變化；淨利＝營業利益 × 0.70（公司指引隱含）", False),
 ("PL｜FY2026 錨點＝公司指引（2026-08-04 上修：營收 189.0、營業利益 44.0、淨利 30.8、EPS 758.83）×（1＋g_anc，基準 0）", False),
 ("Val｜目標價＝0.5 × DCF（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續）＋ 0.5 × Comparable（日本設備同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）；規則值（10 年 P/E 分位 × FY2027 EPS；循環位置 " + CP.fy2027_text(PEAK) + "）為參考", False), ("", False),
 ("會計年度：4 月到隔年 3 月（FY2025＝2025-04～2026-03，公司「2026年3月期」）。金額為日圓，未建匯率（公司前提 1 美元＝150 日圓，日圓計價比重高，匯率影響輕微）。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（PL 頁底部延伸區塊，灰底）。情境：A＝順利、B＝平平、C＝不如預期（設備市場讀 markets/equipment 同字母）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("可識別性：SPE 只有營收與部門營業利益，無產品別金額；探針台／加工裝置比例取自公司「7 割前半／2 割後半」的定性說法；超額成長與中國變數為 Assumed，列在 Sens。", False),
 ("已知限制：見 accretech/README.md「待辦與限制」。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 180

# 設備總額具名範圍（Recon 用）：放在 In 頁之後的隱藏區？——改為寫在 In 頁最後（只讀上游值）
r += 1
IN.cell(r, 1, "K. 設備總額（markets/equipment 成品 Mkt 頁，同字母；只用於 Recon 結構檢查）").font = BOLD; r += 1
for s in SC:
    for j, y in enumerate((25, 26, 27, 28, 29)):
        IN.cell(r, 1, f"tot_{s}_{y}"); IN.cell(r, 2, f"設備總額 CY20{y}（情境 {s}）"); IN.cell(r, 3, "US$B")
        IN.cell(r, 5, EQ[(s, "tot")][0][j]).font = BLUE; IN.cell(r, 7, "Derived"); IN.cell(r, 8, EQ[(s, "tot")][1] if j == 0 else "同上")
        defname(f"tot_{s}_{y}", f"In!$E${r}"); r += 1


# ---------------- v0.1.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00、14:17）；新增列放在各頁現有最後一列之後，既有儲存格不變 ----------------
_U = f"B{T['upw']}"
_G = V.max_row + 2
V.cell(_G, 1, "上檔空間級距")
V[f"B{_G}"] = f'=IF({_U}>0.5,"大於 50%",IF({_U}>0.25,"25%–50%",IF({_U}>=0,"0%–25%",IF({_U}>=-0.25,"0%～−25%",IF({_U}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"B{_G}"].number_format = "@"; V[f"E{_G}"] = "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）"
V.cell(_G + 1, 1, "投資建議")
V[f"B{_G + 1}"] = f'=IF({_U}>0.5,"強力買進",IF({_U}>0.25,"買進",IF({_U}>=-0.25,"中立",IF({_U}>=-0.5,"放空","強力放空"))))'
V[f"B{_G + 1}"].number_format = "@"
V[f"E{_G + 1}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空"
OBS = ["目標價 ¥19,528（DCF 10,189、Comparable 28,867 各半），比現價 15,520（2026-10-23）高 25.8%，級距 25%–50%；Comparable 單獨為 +86.0%。主因是 Comparable 的倍數：同業（Disco、東京威力科創、Advantest、Lasertec、Kokusai Electric）前瞻 P/E 中位數 34.0 倍，東京精密自身前瞻 P/E 只有 21.9 倍、現價隱含 18.3 倍 12 個月後前瞻 EPS（849），Comparable 未扣本公司長期低於同業的折價；EPS 只比共識前瞻 EPS 801 高 5.9%。DCF 反而比現價低 34.3%（ROE 13.9%、再投資率 28.7%）。",
       "共識目標價 20,900 只隱含 26.1 倍（比同業中位數低 23.3%），市場已給折價；若 Comparable 改用 26.1 倍，加權目標價約落在中立區。情境 FY2028 EPS A 1,077、B 831、C 562，現價隱含 P/E 的十年分位 A 44%、B 90%、C 100%。",
       "FY2027 EPS 對 FY2026 錨點（實績低於指引 3% → −3.6%、高於 5% → +6.0%）與中國 SPE（−20% → −4.5%）最敏感，情境 C 為 −18.6%。下一個檢驗事件是 FY2026 第二季財報（約 2026-11 上旬，日期未確認）：上半年實績與全年指引是否維持。"]   # 重點觀察（與 accretech/README.md 相同文字）
_R = RD.max_row + 1
RD.cell(_R, 1, "v0.1.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁最後新增「上檔空間級距」「投資建議」兩列（依「上檔空間：加權目標價」固定規則）；README 頁最後新增投資建議與重點觀察三點。")
RD.cell(_R + 1, 1, f'="投資建議："&Val!B{_G + 1}&"（依上檔空間固定規則）"')
for _i, _t in enumerate(OBS, 1): RD.cell(_R + 1 + _i, 1, f"重點觀察 {_i}：{_t}")
if os.environ.get("G1_ROWS"): print("G1ROWS", "accretech", "grade", _G, "rec", _G + 1, "readme_obs", _R + 2, _R + 4)

if os.environ.get("ROWS"):
    print("T", T, "SUMR", SUMR, "EXT", EXT, "PB", {s: PB[s] for s in SC}, "CHK26", CHK26, "TS0", TS0, "SENS_LAST", SENS_LAST)
order = ["README", "In", "Q", "Hist", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
