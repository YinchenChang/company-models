# Arm Holdings（ARM）個股模型 v0.1（AI 半導體 Project：運算次產業；CPU 架構授權與權利金）
# 用法：python3 arm/build_arm.py raw.xlsx ；再以 LibreOffice headless 重算（見 README）
# 會計年度：FY(t) 於 t 年 3 月 31 日結束（FY2027＝2026-04～2027-03）。上游市場層是日曆年：FY(t) 市場＝9/12 × CY(t−1)＋3/12 × CY(t)。
# 金額：十億美元（US$B）；EPS：美元，non-GAAP（公司與市場共識口徑，排除股票薪酬、併購攤銷；GAAP EPS 另列對照）。
# 樣板：crdo/（美股、估值 v2）與 tuc/（敏感度＝一次改一個 In 頁低／高值、整本重算；arm/sens_arm.py 產生 arm/sens.json）。
# 官方營收：授權及其他（License and other）、權利金（Royalty）兩類（公司單一營運部門，營收依類型揭露）。
# 命題：權利金＝晶片出貨 × 每顆權利金率，分智慧型手機、資料中心（Neoverse）、車用、物聯網與其他；資料中心 NVIDIA Grace／Vera CPU 顆數讀 Siliconomics 加速器顆數。
# 環境變數 ARM_OVERRIDE="鍵=值;…"：只供 sens_arm.py 使用；CI 與正式建置不設。
import sys, os, json
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # noqa: E402
sys.path.insert(0, os.path.join(ROOT, "tools"))
import relative_multiple as RMX   # noqa: E402
sys.path.insert(0, os.path.join(ROOT, "tools", "in_layout"))
import apply_in_layout   # noqa: E402

VER, DATE = "v0.1", "2026-10-31"
BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
SC = ("A", "B", "C")
OVR = {}
for kv in filter(None, os.environ.get("ARM_OVERRIDE", "").split(";")):
    k, v = kv.split("="); OVR[k] = float(v)
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
SNAP = f"Siliconomics 快照 {SIL.VERSION}（{SIL.FILE}）"
N = SIL.NAMES
UP = {k: N[f"SIF_UnitsPlat_{k}"] for k in ("bw", "rb", "ru")}    # [2025(None), 2026, 2027, 2028, 2029]
UC = {s: N[f"SIF_UnitsCoef_{s}"] for s in SC}

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E）。金額：十億美元（US$B）；出貨：十億顆或百萬顆；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = ("標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。公司原文＝Arm 股東信與 6-K（SEC）、20-F、法說；公司指引為 Interested-party；"
            "共識與券商皆為 Interested-party。行業變數（加速器顆數、情境係數）讀 " + SNAP + "；Siliconomics 沒有的行業變數（手機出貨、雲端自研 CPU 顆數）放本頁並標「待併入 Siliconomics」。")
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
SL26 = "https://www.sec.gov/Archives/edgar/data/0001973239/000197323926000062/exhibit992fye26q431-marx26.htm"
SL27 = "https://www.sec.gov/Archives/edgar/data/0001973239/000197323926000113/exhibit992fye27q130-junx26.htm"
SL25 = "https://www.sec.gov/Archives/edgar/data/1973239/000197323925000010/exhibit992fye25q431-marx25.htm"
F20 = "https://www.sec.gov/Archives/edgar/data/0001973239/000197323926000097/arm-20260331.htm"
TR3 = "https://www.fool.com/earnings/call-transcripts/2026/02/04/arm-holdings-arm-q3-2026-earnings-transcript/"
LB = "https://longbridge.com/news/294985711"
HL = "https://www.hudson-labs.com/research/arm-q2-2027-earnings-preview-arm-revenue-and-outlook"
PW = "https://www.paulweiss.com/insights/client-news/qualcomm-wins-jury-verdict-in-breach-of-contract-and-interference-trial-against-arm"
SAF = "https://stockanalysis.com/stocks/arm/forecast/"
IS = "https://www.investsmart.com.au/security/nasdaq/arm/arm-holdings-plc/share-price"

sec("A. FY2027 錨點（1Q27 實績＋2Q27 公司指引＋下半年假設；2Q27 財報 2026-11-04 公布）")
inp("q2_lg", "2Q27 授權及其他年增（公司指引：mid-30%）", "%", 0.30, 0.35, 0.38, "Interested-party", "公司 2Q27 指引：授權年增 mid-30%、權利金 low-teens、營收 13.8 億 ± 0.5 億美元（Longbridge 2026-07-29 轉述股東信）", "Hudson Labs：授權及其他年增約 30%")
inp("q2_rg", "2Q27 權利金年增（公司指引：low-teens）", "%", 0.10, 0.12, 0.14, "Interested-party", "同上", "去年同期 6.20 億美元（2Q26）")
inp("h2_lg", "3Q27、4Q27 授權及其他年增", "%", 0.08, 0.20, 0.30, "Assumed", "ACV 1Q27 年增 13%（股東信）；1Q27 授權年增 23%、2Q27 指引 mid-30%", "去年同期 3Q26 5.05、4Q26 8.19 億美元（含 SoftBank 協議每季約 2 億美元）")
inp("h2_rg", "3Q27、4Q27 權利金年增", "%", 0.08, 0.15, 0.20, "Assumed", "1Q27 +22%、2Q27 指引 low-teens；資料中心權利金年增逾一倍（股東信）", "管理層：FY2027 約 20% 成長「非常合理」（3Q26 法說，The Motley Fool 逐字稿）")
inp("q2_opx", "2Q27 non-GAAP 營業費用", "US$B", None, 0.78, None, "Interested-party", "公司 2Q27 指引約 7.8 億美元（股東信 2026-07-29）", "1Q27 實績 7.33 億")
inp("q3_opx", "3Q27 non-GAAP 營業費用", "US$B", 0.78, 0.80, 0.83, "Assumed", "2Q27 指引 7.8 億；FY2026 季增約 2–4%（股東信）", "AGI CPU 與 CSS 研發投入")
inp("q4_opx", "4Q27 non-GAAP 營業費用", "US$B", 0.80, 0.82, 0.86, "Assumed", "同上", "—")
inp("cogs_r", "IP 營業成本 ÷ 營收（non-GAAP；不含晶片）", "%", 0.015, 0.019, 0.025, "Derived", "1Q27 non-GAAP 毛利率 98.1%（股東信）", "FY2026：（營收 − 營業費用 − 營業利益）÷ 營收＝1.8%")
inp("q_oth", "每季業外收益（利息，non-GAAP）", "US$B", 0.025, 0.034, 0.04, "Derived", "1Q27 non-GAAP 稅前 5.65 億 − 營業利益 5.31 億（股東信）", "現金與短期投資利息")
inp("q_tx", "non-GAAP 有效稅率（FY2027）", "%", 0.13, 0.15, 0.18, "Derived", "1Q27 non-GAAP 稅 0.85 ÷ 稅前 5.65＝15.0%（股東信調節表）", "英國 OECD 最低稅 15%")
inp("q_sh", "稀釋加權股數（FY2027）", "十億股", None, 1.080, None, "Derived", "1Q27 1,078 百萬股（股東信）", "4Q26 1,068；20-F：2026-03-31 普通股 1,064 百萬股")
inp("chip_27", "Arm AGI CPU 晶片營收 FY2027（含在授權及其他錨點內）", "US$B", 0.0, 0.15, 0.40, "Assumed", "公司：AGI CPU 2026-03 推出，首批已交付多家美、中客戶；FY2027–28 客戶需求逾 20 億美元（股東信 2026-07-29）", "晶片營收所在列公司未揭露：未找到；本模型暫列授權及其他之下")

sec("B. 權利金分市場（FY2026 基期比例；公司只揭露手機占比）")
inp("ph_s", "FY2026 智慧型手機（行動應用處理器）÷ 權利金", "%", 0.40, 0.43, 0.46, "Verified", "20-F（FY2026）：行動應用處理器權利金約占權利金 43%", "管理層：手機約 40–45%（3Q26 法說）")
inp("dc_s", "FY2026 資料中心（Neoverse）÷ 權利金", "%", 0.10, 0.14, 0.19, "Interested-party", "管理層：資料中心「十幾到接近 20%」（3Q26 法說，The Motley Fool 逐字稿）", "1Q27 資料中心權利金年增逾一倍（股東信）")
inp("au_s", "FY2026 車用 ÷ 權利金", "%", 0.07, 0.09, 0.11, "Assumed", "管理層：車用雙位數成長（3Q26 法說）；占比未揭露：未找到", "物聯網與其他（含 PC、消費電子、嵌入式）＝殘差")
inp("ph_u26", "FY2026 智慧型手機出貨（全球）", "十億支", 1.15, 1.22, 1.28, "Assumed", "全球手機年出貨約 12 億支（IDC 量級；本專案未建立）", "Siliconomics 無手機出貨：待併入 Siliconomics")
inp("g_ph27", "FY2027 手機權利金年增", "%", 0.08, 0.12, 0.18, "Interested-party", "Hudson Labs：管理層預期 FY2027 手機權利金年增 high-teens（原 20% 以上）", "記憶體漲價使手機出貨下修（聯發科預期約 −15%；管理層：出貨 −20% 使權利金 −1–2%）；本模型取保守值")
inp("g_dc27", "FY2027 資料中心權利金年增", "%", 0.50, 0.70, 0.95, "Interested-party", "1Q27 資料中心權利金年增逾一倍（股東信）；全年假設低於上半年（Assumed）", "Neoverse 核心出貨逾 15 億顆，近 9 個月 5 億（新聞稿 2026-07-29）")
inp("g_au27", "FY2027 車用權利金年增", "%", 0.08, 0.15, 0.22, "Assumed", "管理層：車用雙位數成長（3Q26 法說）", "物聯網與其他 FY2027＝錨點權利金 − 三個市場（殘差）")

sec("C. 權利金驅動（FY2028 起：出貨 × 每顆權利金率）")
for y, v in ((28, (-0.08, -0.02, 0.03)), (29, (-0.03, 0.01, 0.04))):
    inp(f"g_phu{y}", f"手機出貨年增 FY20{y}", "%", *v, "Assumed", "記憶體價格上漲壓抑 2026–27 手機出貨（聯發科預期約 −15%，3Q26 法說轉述）" if y == 28 else "成熟市場，回到低個位數（Assumed）", "Siliconomics 無手機出貨：待併入 Siliconomics")
for s, v in (("A", 0.14), ("B", 0.11), ("C", 0.07)):
    inp(f"phr_{s}", f"情境 {s}：手機每支權利金率年增（FY2028–29；v9、CSS 滲透）", "%", None, v, None, "Assumed",
        "管理層：v9 CSS 權利金率每代上升；CSS 占權利金 FY2026 進入十幾 %、未來兩年可能逾 50%（3Q26 法說）" if s == "A" else ("同上；情境 B 為基準" if s == "B" else "情境 C：Qualcomm 2026-10-13 陪審團判 Arm 違約（架構授權），Arm 難以要求 Qualcomm 改用較高費率；加上 CSS 推廣較慢"),
        "Qualcomm 訴訟：Paul Weiss 2026-10-13；法官審理部分 10-20 結辯，判決未出" if s == "C" else "—")
inp("qc_adj", "Qualcomm 相關手機權利金調整（FY2028 起；基準 0）", "%", -0.06, 0.0, 0.02, "Assumed", "陪審團 2026-10-13 判 Arm 違反架構授權、干擾 Qualcomm 與 Meta 的業務關係（Paul Weiss）；損害賠償未揭露：未找到", "低值＝Qualcomm 擴大自研核心、Arm 失去部分手機權利金；高值＝和解後費率調升")
inp("cpu_gpu", "每顆 NVIDIA 資料中心 GPU 搭配的 Arm CPU（Grace／Vera）", "顆", 0.4, 0.5, 0.5, "Interested-party", "GB200／GB300 NVL72：72 顆 GPU 搭 36 顆 Grace；VR NVL144：72 顆封裝搭 36 顆 Vera（NVIDIA）", "低值反映 HGX（x86 主機）出貨占比")
for k, lab in (("bw", "Blackwell"), ("rb", "Rubin"), ("ru", "Rubin Ultra")):
    for y, j in ((26, 1), (27, 2), (28, 3), (29, 4)):
        inp(f"nv{k}_{y}", f"NVIDIA {lab} 顆數 CY20{y}（情境 A）", "百萬顆", None, float(UP[k][j]), None, "Derived", f"{SNAP}｜SIF_UnitsPlat_{k}[20{y}]", "Siliconomics 加速器顆數（情境 A）；B、C 乘同字母顆數係數")
for y, j in ((28, 3), (29, 4)):
    for s in SC:
        inp(f"uc{y}_{s}", f"情境 {s}：CY20{y} 顆數係數", "x", None, float(UC[s][j]), None, "Derived", f"{SNAP}｜SIF_UnitsCoef_{s}[20{y}]", "A＝1；B＝0.8／0.85；C＝0.6（與台積電、NVIDIA 模型同一套）")
inp("nv_cy25", "NVIDIA 資料中心 GPU 顆數 CY2025（Blackwell＋Hopper）", "百萬顆", 5.0, 6.0, 7.0, "Assumed", "Siliconomics 加速器顆數從 2026 年起：CY2025 未建立，待併入 Siliconomics", "只用於 FY2026 的 NVIDIA CPU 指數基期")
inp("w_nv", "FY2027 資料中心權利金中 NVIDIA Grace／Vera 的比重", "%", 0.20, 0.35, 0.50, "Assumed", "公司未揭露客戶別：未找到；NVIDIA Vera 88 核、AWS Graviton5 192 核、Microsoft Cobalt 200 132 核（3Q26 法說）", "其餘＝雲端自研 CPU（Graviton、Axion、Cobalt 等）與網通、DPU")
inp("g_hc", "雲端自研 Arm CPU 顆數年增（Graviton、Axion、Cobalt 等；× 同字母顆數係數）", "%", 0.15, 0.30, 0.45, "Assumed", "管理層：預期 Arm 在前幾大雲端業者的 CPU 份額達 50%（3Q26 法說）；IDC：Arm 加速伺服器平台支出兩季近倍增（新聞稿 2026-07-29）", "Siliconomics 無 CPU 顆數：待併入 Siliconomics")
inp("dcr", "資料中心每顆權利金率年增（核心數、CSS）", "%", 0.04, 0.10, 0.15, "Assumed", "核心數：Graviton4 96 → Graviton5 192、Vera 72 → 88（3Q26 法說）", "—")
inp("g_au", "車用權利金年增（FY2028 起）", "%", 0.08, 0.15, 0.22, "Assumed", "車用雙位數成長（3Q26 法說）", "—")
inp("g_iot", "物聯網與其他權利金年增（FY2028 起）", "%", -0.02, 0.04, 0.08, "Assumed", "含 PC（Windows on Arm、NVIDIA RTX Spark）、消費電子、嵌入式；FY2027 殘差約持平", "—")

sec("D. 授權及其他（含 SoftBank 協議）與 Arm AGI CPU 晶片")
inp("g_lic", "授權（不含晶片）年增（FY2028 起）", "%", 0.02, 0.10, 0.15, "Assumed", "ACV 1Q27 年增 13%（股東信）；SoftBank 協議 3Q26 認列 2.0 億美元／季（3Q26 法說）", "SoftBank 協議期間與續約：未找到")
for y, vals in ((28, (1.4, 1.0, 0.5)), (29, (3.0, 2.0, 0.8))):
    for s, v in zip(SC, vals):
        inp(f"chip{y}_{s}", f"情境 {s}：Arm AGI CPU 晶片營收 FY20{y}", "US$B", None, v, None, "Assumed",
            "客戶需求 FY2027–28 逾 20 億美元，已取得 10 億美元的產能（股東信 2026-07-29）；公司稱朝此業務 150 億美元的預估前進（4Q26 股東信）" if (y == 28 and s == "A") else "同上",
            "產能為主要限制：第二個 10 億美元的製造產能尚未鎖定（Yahoo Finance 轉述 WSJ）" if (y == 28 and s == "A") else "—")
inp("chip_gm", "晶片毛利率", "%", 0.35, 0.45, 0.55, "Assumed", "晶片毛利率公司未揭露：未找到", "Analogy：CPU 晶片業者 50–60%；首代晶片、委外製造與客戶議價（Assumed 較低）")

sec("E. 營業費用、業外、稅與股數（FY2028 起；non-GAAP）")
inp("g_opx", "non-GAAP 營業費用年增", "%", 0.08, 0.13, 0.18, "Assumed", "FY2026 non-GAAP 營業費用年增 33%（研發 +43%）；FY2027 錨點約 +15%（股東信）", "晶片業務需持續研發與銷售投入")
inp("oth_f", "業外收益（每年）", "US$B", 0.08, 0.14, 0.18, "Derived", "1Q27 0.034 × 4（股東信）", "—")
inp("tx_f", "non-GAAP 有效稅率", "%", 0.13, 0.15, 0.18, "Derived", "同 q_tx", "—")
inp("sh_g", "稀釋股數年變動", "%", 0.0, 0.008, 0.015, "Assumed", "FY2025 1,063 → FY2026 1,068 → 1Q27 1,078 百萬股（股東信）", "股票薪酬發行；公司未大量買回")

sec("F. 共識與券商（只作 Recon 對帳，不作校準）")
inp("c_eps27", "共識 FY2027 non-GAAP EPS（S&P Global 平均）", "US$", None, 2.23, None, "Interested-party", SAF + "（stockanalysis，更新 2026-10-21）", "高 2.38、低 2.06")
inp("c_rev27", "共識 FY2027 營收（平均）", "US$B", None, 6.1, None, "Interested-party", SAF, "Hudson Labs：60.58 億美元")
inp("c_q2rev", "共識 2Q27 營收", "US$B", None, 1.376, None, "Interested-party", HL, "公司指引 13.8 億 ± 0.5 億")

sec("G. 價格與股權")
inp("px", "股價（2026-10-30 收盤）", "US$", None, 228.88, None, "Interested-party", "InvestSmart 價格歷史：2026-10-30 收盤 228.88（+0.77%）", "digrin 10 月現價 228.77；stockanalysis 10-23 收盤 210.33")
inp("sb_own", "SoftBank 集團持股（約）", "%", None, 0.90, None, "Interested-party", "20-F：SoftBank 集團為控股股東；約九成（whoistheownerof「SoftBank's 90% grip」）", "流通股約一成：只作說明，不進目標價")

sec("H. 目標價（DCF 與 Comparable 加權）與收斂法參數（黃底＝Assumed，Andy 可改）")
for s_, w_, nm_ in (("A", 0.3, "順利"), ("B", 0.5, "平平"), ("C", 0.2, "不如預期")):
    inp(f"tp_w{s_}", f"目標價情境權重：情境 {s_}（{nm_}）", "權重", None, w_, None, "Assumed", "Andy 可改（2026-10-25 目標價規格預設 A／B／C＝0.3／0.5／0.2）", "三個權重合計應為 1")
inp("tp_r", "折現率／要求報酬（DCF）", "%", 0.09, 0.10, 0.11, "Assumed", "Andy 可改（工作單預設 10%）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（DCF 永續；情境 B 的 FY2032 營收年增收斂目標）", "%", 0.03, 0.04, 0.05, "Assumed", "Andy 可改（工作單預設 4%）", "Sens 頁 ±1pt（同時移動 PL 收斂目標 fade_g）")
inp("tp_roe", "ROE（近三年平均）", "%", None, "=AVERAGE(0.0655,0.1305,0.1195)", None, "Derived", "stockanalysis 比率頁（Return on Equity，GAAP）：FY2024 6.55%、FY2025 13.05%、FY2026 11.95%", "單一來源；GAAP 淨利口徑（non-GAAP ROE 較高，見 README）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", 0.3, 0.5, 0.7, "Assumed", "Andy 可改（2026-10-25 11:53：DCF／Comparable 各半）", "Sens 頁 0.3／0.7")
inp("fade_gA", "情境 A（順利）：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂，FY2032 達到")
inp("fade_gB", "情境 B（平平）：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C（不如預期）：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("nim10_f", "十年平均淨利率的替代：最後一年（FY2029）淨利率 ×（本值）", "x", None, 0.8, None, "Assumed", "十年平均淨利率：未找到（2023 年 9 月才上市）；依工作單以最後一年 × 0.8 替代", "淨利率收斂目標：A＝FY2029、B＝中點、C＝替代值")
inp("cons_tp", "市場共識目標價（賣方平均）", "US$", None, 290.71, None, "Interested-party", "stockanalysis forecast 頁（S&P Global）43 位分析師平均 290.71（中位數 278.60），更新 2026-10-21", "UBS 320、RBC 340（marketscreener，2026-07-30）")
inp("cons_feps", "共識前瞻 12 月 EPS（＝價格 ÷ 前瞻 P/E）", "US$", None, "=225.95/93.52", None, "Interested-party", "stockanalysis 個股頁（2026-10-27）：價格 225.95 ÷ Forward PE 93.52", "賣方共識，只作對帳與差距拆解，不作校準")

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    if key in OVR: base = OVR[key]
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    defname(key, f"In!$E${r}"); INROW[key] = r
    if key.startswith(("tp_", "fade_g", "nim10")): IN.cell(r, 5).fill = KEY
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
INLIST = [x for x in rows if x[0] != "sec"]

# ---------------- Q：實績與 FY2027 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2027 錨點（十億美元；EPS 美元；股數十億股；non-GAAP）。藍字＝輸入；黑字＝公式。實績取自 Arm 股東信（6-K Exhibit 99.2，SEC）"
Q["A2"] = "FY2027 錨點＝1Q27 實績（2026-07-29）＋2Q27 公司指引（授權 mid-30%、權利金 low-teens、營業費用 7.8 億）＋3Q27、4Q27 假設（以去年同期 × (1＋年增)）。2Q27 財報 2026-11-04（美東盤後）公布後更新：待 FQ2 財報"
hdr(Q, 3, ["項目", "FY2025A", "FY2026A", "1Q27（實績）", "2Q27（指引）", "3Q27（模型）", "4Q27（模型）", "FY2027 錨點", "標記", "來源1", "來源2／說明"])
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1="", s2="", ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"H{r}"] = ann or f"=SUM(D{r}:G{r})"; Q[f"H{r}"].number_format = fmt
    Q.cell(r, 9, tag); Q.cell(r, 10, s1); Q.cell(r, 11, s2)
qrow(4, "授權及其他營收（License and other）", {"B": 1.839, "C": 2.307, "D": 0.574, "E": "=C21*(1+q2_lg)", "F": "=C22*(1+h2_lg)", "G": "=C23*(1+h2_lg)"}, tag="Verified／指引／Assumed",
     s1="FY2025：" + SL25 + "；FY2026：" + SL26, s2="1Q27：" + SL27)
qrow(5, "權利金營收（Royalty）", {"B": 2.168, "C": 2.613, "D": 0.715, "E": "=D21*(1+q2_rg)", "F": "=D22*(1+h2_rg)", "G": "=D23*(1+h2_rg)"}, tag="Verified／指引／Assumed", s1="同上", s2="同上")
qrow(6, "營收", {c: f"={c}4+{c}5" for c in "BCDEFG"}, tag="Derived", s1="FY2025 40.07、FY2026 49.20、1Q27 12.89 億美元（股東信）", s2="2Q27 指引 13.8 億 ± 0.5 億")
qrow(7, "營業成本（non-GAAP）", {"B": "=B6-B8-B9", "C": "=C6-C8-C9", "D": "=D6-D8-D9", "E": "=E6*cogs_r", "F": "=F6*cogs_r", "G": "=G6*cogs_r"}, tag="Derived", s1="營收 − non-GAAP 營業費用 − non-GAAP 營業利益", s2="1Q27 non-GAAP 毛利率 98.1%")
qrow(8, "營業費用（non-GAAP；研發＋銷管）", {"B": "=1.340+0.709", "C": "=1.911+0.806", "D": "=0.530+0.203", "E": "=q2_opx", "F": "=q3_opx", "G": "=q4_opx"}, tag="Verified／指引／Assumed", s1="股東信 non-GAAP 調節表", s2="")
qrow(9, "營業利益（non-GAAP）", {"B": 1.871, "C": 2.115, "D": 0.531, "E": "=E6-E7-E8", "F": "=F6-F7-F8", "G": "=G6-G7-G8"}, tag="Verified／Derived", s1="股東信", s2="1Q27 營益率 41.2%")
qrow(10, "營業利益率（non-GAAP）", {c: f"={c}9/{c}6" for c in "BCDEFGH"}, fmt="0.0%", ann="", tag="Derived")
qrow(11, "業外收益（non-GAAP）", {"D": "=0.565-0.531", "E": "=q_oth", "F": "=q_oth", "G": "=q_oth"}, tag="Derived／Assumed", s1="1Q27 稅前 5.65 − 營業利益 5.31（股東信）", s2="FY2025、FY2026 含在淨利內")
qrow(12, "稅前利益（non-GAAP）", {"D": 0.565, "E": "=E9+E11", "F": "=F9+F11", "G": "=G9+G11"}, tag="Verified／Derived", s1="股東信", s2="")
qrow(13, "稅後淨利（non-GAAP）", {"B": 1.737, "C": 1.889, "D": 0.480, "E": "=E12*(1-q_tx)", "F": "=F12*(1-q_tx)", "G": "=G12*(1-q_tx)"}, tag="Verified／Derived", s1="股東信", s2="1Q27 稅 0.85 億（15%）")
qrow(14, "稀釋加權股數（十億股）", {"B": 1.063, "C": 1.068, "D": 1.078, "E": "=q_sh", "F": "=q_sh", "G": "=q_sh"}, ann="=q_sh", tag="Verified／Derived", s1="股東信", s2="")
qrow(15, "EPS（non-GAAP）", {c: f"={c}13/{c}14" for c in "BCDEFG"}, fmt="0.00", ann="=H13/H14", tag="Derived", s1="淨利 ÷ 股數", s2="")
qrow(16, "  公司揭露 non-GAAP EPS（2Q27 為指引中點）", {"B": 1.63, "C": 1.77, "D": 0.45, "E": 0.47}, fmt="0.00", ann="", tag="Verified／Interested-party", s1="股東信", s2="2Q27 指引 0.47 ± 0.04")
qrow(17, "  公司揭露 GAAP EPS（對照；含股票薪酬）", {"B": 0.75, "C": 0.85, "D": 0.25}, fmt="0.00", ann="", tag="Verified", s1="股東信", s2="1Q27 股票薪酬 3.43 億美元（股東信）")
Q.cell(19, 1, "去年同期季資料（FY2026，十億美元；3Q27、4Q27 年增的基期）").font = BOLD
hdr(Q, 20, ["季", "", "授權及其他", "權利金", "營收"])
QR = {}
for i, (q_, l_, r_) in enumerate((("2Q26", 0.515, 0.620), ("3Q26", 0.505, 0.737), ("4Q26", 0.819, 0.671), ("1Q26", 0.468, 0.585))):
    rr = 21 + i; QR[q_] = rr
    Q.cell(rr, 1, q_); Q.cell(rr, 3, l_).font = BLUE; Q.cell(rr, 4, r_).font = BLUE; Q.cell(rr, 5, f"=C{rr}+D{rr}")
    Q.cell(rr, 9, "Verified" if q_ in ("1Q26", "4Q26") else "Derived")
    Q.cell(rr, 10, {"1Q26": "1Q27 股東信（去年同期欄）", "4Q26": "4Q26 股東信", "3Q26": "3Q26 法說：權利金 7.37、授權 5.05 億（The Motley Fool 逐字稿）", "2Q26": "＝FY2026 − 1Q26 − 3Q26 − 4Q26（股東信全年數）"}[q_])
Q["C21"] = "=C4-C22-C23-C24"; Q["D21"] = "=C5-D22-D23-D24"; Q["C21"].font = Font(); Q["D21"].font = Font()
Q.cell(26, 1, "核對：四季合計 − FY2026（應為 0）"); Q["C26"] = "=SUM(C21:C24)-C4"; Q["D26"] = "=SUM(D21:D24)-C5"
Q.cell(28, 1, "FY2027 錨點的檢查（參考；校準未使用共識）").font = BOLD
Q.cell(29, 1, "FY2027 營收年增"); Q["B29"] = "=H6/C6-1"
Q.cell(30, 1, "FY2027 營收 ÷ 共識 − 1"); Q["B30"] = "=H6/c_rev27-1"
Q.cell(31, 1, "FY2027 EPS ÷ 共識 − 1"); Q["B31"] = "=H15/c_eps27-1"
Q.cell(32, 1, "2Q27 營收（模型）÷ 公司指引中點 − 1"); Q["B32"] = "=E6/1.38-1"
Q.cell(33, 1, "2Q27 EPS（模型）÷ 公司指引中點 − 1"); Q["B33"] = "=E15/0.47-1"
Q.cell(34, 1, "FY2027 營業費用年增"); Q["B34"] = "=H8/C8-1"
for rr in range(29, 35): Q[f"B{rr}"].number_format = "0.0%"
for nm, rr in (("lo", 4), ("ro", 5), ("rev", 6), ("cogs", 7), ("opx", 8), ("oi", 9), ("ni", 13), ("sh", 14), ("eps", 15)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$C${rr}"); defname(f"a27_{nm}", f"Q!$H${rr}")
Q.column_dimensions["A"].width = 58
for c in "BCDEFGH": Q.column_dimensions[c].width = 13
Q.column_dimensions["I"].width = 22; Q.column_dimensions["J"].width = 80; Q.column_dimensions["K"].width = 60
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
cyu = lambda y, s: f"(nvbw_{y}+nvrb_{y}+nvru_{y})" + (f"*uc{y}_{s}" if y in (28, 29) else "")   # CY 顆數（情境 s）
cyc = lambda y, s: (f"uc{y}_{s}" if y in (28, 29) else "1")
def engine(ws, r0, title, s):
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.000", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    sec_("A. 資料中心 CPU 顆數指數（NVIDIA Grace／Vera 讀 Siliconomics；雲端自研 CPU Assumed × 同字母係數）")
    row("nvcy", "NVIDIA 資料中心 GPU 顆數（會計年度＝9/12 × CY(t−1)＋3/12 × CY(t)）", "百萬顆",
        {"D": f"=0.75*nv_cy25+0.25*{cyu(26, s)}", "E": f"=0.75*{cyu(26, s)}+0.25*{cyu(27, s)}", "F": f"=0.75*{cyu(27, s)}+0.25*{cyu(28, s)}", "G": f"=0.75*{cyu(28, s)}+0.25*{cyu(29, s)}"},
        fmt="0.00", note=f"Siliconomics SIF_UnitsPlat_bw＋rb＋ru × SIF_UnitsCoef_{s}；CY2025 為 In 頁 nv_cy25（待併入 Siliconomics）")
    row("nvcpu", "NVIDIA Grace／Vera CPU 顆數 ＝ GPU × 每顆 GPU 搭配 CPU", "百萬顆", {c: f"={c}{R['nvcy']}*cpu_gpu" for c in "DEFG"}, fmt="0.00")
    row("hcx", "雲端自研 Arm CPU 顆數指數（FY2027＝1）", "x", {"E": "=1", "F": f"=E{r}*(1+g_hc)*(0.75*{cyc(27, s)}+0.25*{cyc(28, s)})", "G": f"=F{r}*(1+g_hc)*(0.75*{cyc(28, s)}+0.25*{cyc(29, s)})/(0.75*{cyc(27, s)}+0.25*{cyc(28, s)})"},
        fmt="0.00", note="年增 g_hc × 同字母顆數係數（行業情境）；Siliconomics 無 CPU 顆數：待併入 Siliconomics")
    row("dcx", "資料中心 CPU 顆數指數（FY2027＝1）＝ w_nv × NVIDIA 指數＋（1 − w_nv）× 雲端自研指數", "x",
        {"E": "=1", **{c: f"=w_nv*{c}{R['nvcpu']}/$E${R['nvcpu']}+(1-w_nv)*{c}{R['hcx']}" for c in "FG"}}, fmt="0.00")
    row("dcrx", "資料中心每顆權利金率指數（FY2027＝1）", "x", {"E": "=1", "F": "=E{0}*(1+dcr)".format(r), "G": "=F{0}*(1+dcr)".format(r)}, fmt="0.00")
    sec_("B. 權利金（分市場）＝ 出貨 × 每顆權利金率")
    row("phu", "智慧型手機出貨", "十億支", {"D": "=ph_u26", "E": f"=D{r}", "F": f"=E{r}*(1+g_phu28)", "G": f"=F{r}*(1+g_phu29)"}, fmt="0.000",
        note="FY2027 出貨假設持平（權利金成長來自費率）；Siliconomics 無手機出貨：待併入 Siliconomics")
    row("phrt", "手機每支權利金（美元）＝ 手機權利金 ÷ 出貨", "US$", {"D": f"=ph_s*a26_ro/D{R['phu']}", "E": f"=D{r}*(1+g_ph27)", "F": f"=E{r}*(1+phr_{s})", "G": f"=F{r}*(1+phr_{s})"}, fmt="0.000",
        note=f"FY2026 Derived；FY2028–29 年增＝情境 {s} 的 phr_{s}（v9、CSS 滲透；C 含 Qualcomm 判決影響）")
    row("ph", "權利金：智慧型手機（行動應用處理器）", "US$B", {"C": f"=ph_s*a25_ro", "D": f"=ph_s*a26_ro", "E": f"=E{R['phu']}*E{R['phrt']}",
        **{c: f"={c}{R['phu']}*{c}{R['phrt']}*(1+qc_adj)" for c in "FG"}}, note="FY2025 拆分沿用 FY2026 比例（Assumed）；FY2028 起 × (1＋qc_adj)")
    row("dc", "權利金：資料中心（Neoverse；Grace、Vera、Graviton、Axion、Cobalt 等）", "US$B",
        {"C": f"=dc_s*a25_ro", "D": f"=dc_s*a26_ro", "E": f"=D{r}*(1+g_dc27)", **{c: f"=$E{r}*{c}{R['dcx']}*{c}{R['dcrx']}" for c in "FG"}}, note="FY2028 起＝FY2027 × CPU 顆數指數 × 每顆權利金率指數")
    row("au", "權利金：車用", "US$B", {"C": f"=au_s*a25_ro", "D": f"=au_s*a26_ro", "E": f"=D{r}*(1+g_au27)", "F": f"=E{r}*(1+g_au)", "G": f"=F{r}*(1+g_au)"})
    row("iot", "權利金：物聯網與其他（含 PC、消費電子、嵌入式；FY2027 前為殘差）", "US$B",
        {"C": f"=a25_ro-C{R['ph']}-C{R['dc']}-C{R['au']}", "D": f"=a26_ro-D{R['ph']}-D{R['dc']}-D{R['au']}", "E": f"=a27_ro-E{R['ph']}-E{R['dc']}-E{R['au']}", "F": f"=E{r}*(1+g_iot)", "G": f"=F{r}*(1+g_iot)"},
        note="FY2025–FY2027 為權利金 − 三個市場（殘差）；FY2028 起依成長率")
    row("ro", "權利金營收（Royalty）", "US$B", {c: f"={c}{R['ph']}+{c}{R['dc']}+{c}{R['au']}+{c}{R['iot']}" for c in "CDEFG"}, note="官方營收類型；FY2025–FY2027 等於財報／錨點")
    row("rog", "  權利金年增", "%", {c: f"={c}{R['ro']}/{p}{R['ro']}-1" for p, c in zip("CDEF", "DEFG")}, fmt="0.0%")
    sec_("C. 授權及其他（含 SoftBank 協議）與 Arm AGI CPU 晶片")
    row("chip", "授權及其他：Arm AGI CPU 晶片營收（公司未揭露所在列，暫列於此）", "US$B", {"C": "=0", "D": "=0", "E": "=chip_27", "F": f"=chip28_{s}", "G": f"=chip29_{s}"},
        note=f"情境 {s}：chip28_{s}、chip29_{s}（Assumed；A 最高、C 最低）")
    row("lic", "授權及其他：授權（不含晶片；含 SoftBank 協議）", "US$B", {"C": "=a25_lo", "D": "=a26_lo", "E": f"=a27_lo-E{R['chip']}", "F": f"=E{r}*(1+g_lic)", "G": f"=F{r}*(1+g_lic)"})
    row("lo", "授權及其他營收（License and other）", "US$B", {c: f"={c}{R['lic']}+{c}{R['chip']}" for c in "CDEFG"}, note="官方營收類型；FY2025–FY2027 等於財報／錨點")
    row("rev", "營收", "US$B", {c: f"={c}{R['lo']}+{c}{R['ro']}" for c in "CDEFG"})
    row("g", "營收年增", "%", {c: f"={c}{R['rev']}/{p}{R['rev']}-1" for p, c in zip("CDEF", "DEFG")}, fmt="0.0%")
    row("segchk", "檢查：分部合計 − 合併營收", "US$B", {c: f"={c}{R['lo']}+{c}{R['ro']}-{c}{R['rev']}+({c}{R['ph']}+{c}{R['dc']}+{c}{R['au']}+{c}{R['iot']}-{c}{R['ro']})+({c}{R['lic']}+{c}{R['chip']}-{c}{R['lo']})" for c in "CDEFG"},
        fmt="0.000", note="授權及其他＋權利金 − 營收，加上兩類營收的子列合計 − 母列；FY2025–FY2027 營收＝財報／錨點（Q 頁）")
    row("segchk2", "核對：營收 − 財報（FY2025、FY2026）／錨點（FY2027）", "US$B", {"C": f"=C{R['rev']}-a25_rev", "D": f"=D{R['rev']}-a26_rev", "E": f"=E{R['rev']}-a27_rev"}, fmt="0.000")
    row("dcsh", "資料中心 ÷ 權利金", "%", {c: f"={c}{R['dc']}/{c}{R['ro']}" for c in "CDEFG"}, fmt="0.0%")
    row("phsh", "智慧型手機 ÷ 權利金", "%", {c: f"={c}{R['ph']}/{c}{R['ro']}" for c in "CDEFG"}, fmt="0.0%")
    sec_("D. 損益（non-GAAP）")
    row("cogs", "營業成本（IP × cogs_r＋晶片 ×（1 − 晶片毛利率））", "US$B", {"C": "=a25_cogs", "D": "=a26_cogs", "E": f"=a27_cogs+E{R['chip']}*(1-chip_gm-cogs_r)",
        **{c: f"=({c}{R['rev']}-{c}{R['chip']})*cogs_r+{c}{R['chip']}*(1-chip_gm)" for c in "FG"}}, note="FY2027 錨點的成本以 cogs_r 計，晶片部分另加（毛利率 chip_gm）")
    row("gp", "毛利", "US$B", {c: f"={c}{R['rev']}-{c}{R['cogs']}" for c in "CDEFG"})
    row("gm", "毛利率", "%", {c: f"={c}{R['gp']}/{c}{R['rev']}" for c in "CDEFG"}, fmt="0.0%")
    row("opx", "營業費用（non-GAAP）", "US$B", {"C": "=a25_opx", "D": "=a26_opx", "E": "=a27_opx", "F": f"=E{r}*(1+g_opx)", "G": f"=F{r}*(1+g_opx)"})
    row("oi", "營業利益（non-GAAP）", "US$B", {c: f"={c}{R['gp']}-{c}{R['opx']}" for c in "CDEFG"})
    row("oim", "營業利益率", "%", {c: f"={c}{R['oi']}/{c}{R['rev']}" for c in "CDEFG"}, fmt="0.0%")
    row("ni", "稅後淨利（non-GAAP）", "US$B", {"C": "=a25_ni", "D": "=a26_ni", "E": f"=a27_ni-E{R['chip']}*(1-chip_gm-cogs_r)*(1-q_tx)", **{c: f"=({c}{R['oi']}+oth_f)*(1-tx_f)" for c in "FG"}},
        note="FY2027＝錨點淨利扣除晶片成本差（晶片含在錨點營收內）")
    row("nim", "淨利率", "%", {c: f"={c}{R['ni']}/{c}{R['rev']}" for c in "CDEFG"}, fmt="0.0%")
    row("sh_", "稀釋股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", "E": "=a27_sh", "F": f"=E{r}*(1+sh_g)", "G": f"=F{r}*(1+sh_g)"}, fmt="0.0000")
    row("eps", "EPS（non-GAAP）", "US$", {c: f"={c}{R['ni']}/{c}{R['sh_']}" for c in "CDEFG"}, fmt="0.00")
    row("epsg", "EPS 年增", "%", {c: f"={c}{R['eps']}/{p}{R['eps']}-1" for p, c in zip("CDEF", "DEFG")}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Arm 損益（十億美元；EPS 美元；non-GAAP）。FY2025A、FY2026A 實績，FY2027 錨點（1Q27 實績＋2Q27 指引＋下半年假設），FY2028 起為模型。FY(t)＝t 年 3 月底結束。三情境各一個完整推導區塊（本頁下方），頂端為摘要；最下方延伸到 FY2032（FY2030 起為收斂法）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：順利（Siliconomics 情境 A 顆數；晶片與手機費率最有利）", "B": "情境 B：平平（Siliconomics 情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85）", "C": "情境 C：不如預期（Siliconomics 情境 C：CY2028–29 ×0.6；Qualcomm 判決壓抑手機費率、晶片放量慢）"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], s); r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "US$B", "#,##0.000"), ("g", "  營收年增", "%", "0.0%"), ("ro", "  權利金", "US$B", "#,##0.000"), ("dc", "    其中資料中心", "US$B", "#,##0.000"),
                                ("lo", "  授權及其他", "US$B", "#,##0.000"), ("chip", "    其中 AGI CPU 晶片", "US$B", "#,##0.000"), ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"),
                                ("ni", "稅後淨利（non-GAAP）", "US$B", "#,##0.000"), ("eps", "EPS（non-GAAP）", "US$", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：營收、權利金、授權及其他 − 財報（FY2026）與錨點（FY2027）").font = BOLD; sr += 1
CHK = sr
for lab, f in (("  情境 A：FY2026 營收差", f"=D{PB['A']['rev']}-a26_rev"), ("  情境 A：FY2027 營收差（錨點）", f"=E{PB['A']['rev']}-a27_rev"),
               ("  情境 A：FY2027 權利金差（錨點）", f"=E{PB['A']['ro']}-a27_ro"), ("  情境 A：FY2027 授權及其他差（錨點）", f"=E{PB['A']['lo']}-a27_lo")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = f; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
assert sr < 60, sr
PL.column_dimensions["A"].width = 70; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

XY = (2025, 2026, 2027, 2028, 2029, 2030, 2031, 2032); XC = {y: "CDEFGHIJ"[i] for i, y in enumerate(XY)}; FADE = (2030, 2031, 2032)
xr = PL.max_row + 3; XT = xr
PL.cell(xr, 1, "延伸：三情境到 FY2032。FY2025–FY2029 引用上方推導區塊；FY2030 起為收斂法（灰底），非驅動因子推導：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 線性收斂到收斂目標，FY2032 達到").font = BOLD
xr += 1; XR = {}; XH = None
for s in SC:
    E = PB[s]; R = {}
    PL.cell(xr, 1, f"延伸｜{SCN_T[s]}").font = BOLD; xr += 1
    hdr(PL, xr, ["項目", "單位", "FY2025A", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    XH = XH or xr; xr += 1
    for key, lab, unit in (("rev", "營收（延伸）", "US$B"), ("g", "  營收年增", "%"), ("ni", "淨利（延伸）", "US$B"), ("nim", "  淨利率", "%"),
                           ("sh", "股數（延伸）", "十億股"), ("eps", "EPS（延伸）", "US$"), ("epsg", "  EPS 年增", "%"),
                           ("tg", "收斂目標：FY2032 營收年增", "%"), ("tm", "收斂目標：FY2032 淨利率", "%")):
        PL.cell(xr, 1, lab); PL.cell(xr, 2, unit); R[key] = xr; xr += 1
    G = lambda k: f"$G${R[k]}"
    for y in XY:
        c = XC[y]; pc = XC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={c}{E['rev']}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            PL[f"{c}{R['ni']}"] = f"={c}{E['ni']}"; PL[f"{c}{R['sh']}"] = f"={c}{E['sh_']}"; PL[f"{c}{R['eps']}"] = f"={c}{E['eps']}"
            PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"={G('g')}+($C${R['tg']}-{G('g')})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"={G('nim')}+($C${R['tm']}-{G('nim')})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['sh']}"] = f"={pc}{R['sh']}*(1+sh_g)"
            PL[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh']}"
        if y > 2025: PL[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("rev", "#,##0.000"), ("g", "0.0%"), ("ni", "#,##0.000"), ("nim", "0.0%"), ("sh", "0.0000"), ("eps", "0.00"), ("epsg", "0.0%")):
            PL[f"{c}{R[key]}"].number_format = fm
            if y in FADE: PL[f"{c}{R[key]}"].fill = FADEFILL
    PL[f"C{R['tg']}"] = f"=fade_g{s}"; PL[f"C{R['tg']}"].number_format = "0.0%"
    PL[f"C{R['tm']}"] = {"A": f"=G{R['nim']}", "B": f"=G{R['nim']}*(1+nim10_f)/2", "C": f"=G{R['nim']}*nim10_f"}[s]; PL[f"C{R['tm']}"].number_format = "0.0%"
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導（上游市場層只到 CY2029）")
    PL.cell(R["tg"], 11, {"A": "In 頁 fade_gA＝tp_g＋2pt", "B": "In 頁 fade_gB＝tp_g", "C": "In 頁 fade_gC＝tp_g−2pt"}[s])
    PL.cell(R["tm"], 11, {"A": "＝本情境 FY2029 淨利率", "B": "＝FY2029 淨利率與十年平均替代值（FY2029 × nim10_f）的中點", "C": "＝十年平均替代值（FY2029 × nim10_f；十年平均未找到）"}[s])
    PL.cell(R["sh"], 11, "FY2030 起股數年變動＝sh_g")
    XR[s] = R; xr += 1
PL.column_dimensions["I"].width = 14; PL.column_dimensions["J"].width = 14; PL.column_dimensions["K"].width = 60

# ---------------- Units：權利金分市場（情境 B）與資料中心 CPU ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 權利金＝出貨 × 每顆權利金率（分市場，情境 B）與資料中心 CPU 顆數（NVIDIA Grace／Vera 讀 Siliconomics 加速器顆數）。會計年度 3 月底"
B_ = PB["B"]
hdr(UN, 3, ["項目", "單位", "FY2026A", "FY2027（錨點）", "FY2028", "FY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, ek, fmt, note in (
        ("nvcy", "NVIDIA 資料中心 GPU 顆數（會計年度加權）", "百萬顆", "nvcy", "0.00", "Siliconomics（情境 B）"),
        ("nvcpu", "NVIDIA Grace／Vera CPU 顆數", "百萬顆", "nvcpu", "0.00", "× 每顆 GPU 搭 0.5 顆 CPU"),
        ("dcx", "資料中心 CPU 顆數指數（FY2027＝1）", "x", "dcx", "0.00", "NVIDIA 與雲端自研加權"),
        ("dcrx", "資料中心每顆權利金率指數", "x", "dcrx", "0.00", "Assumed"),
        ("dc", "資料中心權利金", "US$B", "dc", "#,##0.000", ""),
        ("phu", "智慧型手機出貨", "十億支", "phu", "0.000", "待併入 Siliconomics"),
        ("phrt", "手機每支權利金", "US$", "phrt", "0.000", "Derived（FY2026）；之後 × 費率年增"),
        ("ph", "智慧型手機權利金", "US$B", "ph", "#,##0.000", ""),
        ("au", "車用權利金", "US$B", "au", "#,##0.000", ""),
        ("iot", "物聯網與其他權利金", "US$B", "iot", "#,##0.000", ""),
        ("ro", "權利金合計", "US$B", "ro", "#,##0.000", ""),
        ("dcsh", "資料中心 ÷ 權利金", "%", "dcsh", "0.0%", ""),
        ("chip", "Arm AGI CPU 晶片營收", "US$B", "chip", "#,##0.000", "列於授權及其他")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit); UR[key] = ur
    for c_u, c_p in zip("CDEF", "DEFG"):
        UN[f"{c_u}{ur}"] = f"=IF(ISBLANK(PL!{c_p}{B_[ek]}),\"\",PL!{c_p}{B_[ek]})" if ek in ("dcx", "dcrx", "nvcy", "nvcpu") else f"=PL!{c_p}{B_[ek]}"
        UN[f"{c_u}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); ur += 1
ur += 1
UN.cell(ur, 1, "讀法：權利金成長拆成「出貨」與「每顆權利金率」。手機出貨持平到下滑，成長幾乎全靠 v9 與 CSS 費率；資料中心同時有顆數（NVIDIA Grace／Vera 讀 Siliconomics、雲端自研 CPU 年增 Assumed）與每顆費率（核心數）兩個槓桿。").font = BOLD
UNOTE = ur
UN.column_dimensions["A"].width = 60; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 50

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — 目標價（第 0 節：DCF 與 Comparable 加權）、現價隱含倍數與歷史 P/E 分位檢查（美元；non-GAAP EPS）。本版未做先驗，無 Andy 合理 P/E 對照"
V["A2"] = "三情境欄 B／C／D＝A 順利／B 平平／C 不如預期（讀 Siliconomics 同字母情境）。現價＝In 頁 px（2026-10-30 收盤）。SoftBank 持股約九成、流通股約一成：Comparable 不另作流通量調整（見方法列）"
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
PXF = "#,##0.00"
t0 = 4
V.cell(t0, 1, "0. 目標價（DCF 與 Comparable 加權）").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["h1"] = t0 + 2
r_ = t0 + 3
V.cell(r_, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{r_}"] = "PL 頁延伸表（non-GAAP）；FY2030 起為收斂法"; T["eps_h"] = r_; r_ += 1
EPSR = {}
for y in YRS6:
    tset(r_, f"eps{y}", f"  FY{y}", [f"=PL!{XC[y]}{XR[s]['eps']}" for s in SC], "0.00"); EPSR[y] = r_; r_ += 1
V.cell(r_, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = r_; r_ += 1
tset(r_, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe，GAAP）"); RV = r_; r_ += 1
CFR = {}
for y in YRS6:
    tset(r_, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], "0.00"); CFR[y] = r_; r_ += 1
c1, c6 = CFR[2027], CFR[2032]
tset(r_, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], PXF,
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；基準日 2026 年底，t＝1 對 FY2027（3 月底結束；近似）；低於 0 時取 0"); r_ += 1
tset(r_, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{T['dcfv']}*(1+tp_r)" for c in "BCD"], PXF); r_ += 1
tset(r_, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{T['dcf12']}+tp_wB*C{T['dcf12']}+tp_wC*D{T['dcf12']}", None, None], PXF, "Σ tp_wA／wB／wC × 情境價值 ×(1＋tp_r)；權重見 In 頁"); r_ += 2
V.cell(r_, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; r_ += 1
hdr(V, r_, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["ph"] = r_; r_ += 1
SAS = "stockanalysis 個股頁 Forward PE（分析師共識；價格日期如左）"
PEERS = [("高通（Qualcomm）", "QCOM", 19.47, "2026-10-27", SAS), ("博通（Broadcom）", "AVGO", 18.52, "2026-10-23", "stockanalysis（經 mrvl、amd 模型 2026-10-23 取值；本次頁面快取過舊）"),
         ("邁威爾（Marvell）", "MRVL", 42.96, "2026-10-27", SAS), ("Rambus", "RMBS", 24.24, "2026-10-26", SAS),
         ("新思科技（Synopsys）", "SNPS", 25.72, "2026-10-26", SAS), ("益華電腦（Cadence）", "CDNS", 36.00, "2026-10-27", SAS)]
P1 = r_
for nm, tk, pe, dt, src in PEERS:
    V.cell(r_, 1, nm); V.cell(r_, 2, tk); cc = V.cell(r_, 3, pe); cc.font = BLUE; cc.number_format = "0.0"
    V.cell(r_, 4, dt); V.cell(r_, 5, src); V.cell(r_, 6, "Interested-party（賣方共識）"); r_ += 1
P2 = r_ - 1; T["p1"], T["p2"] = P1, P2
tset(r_, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{P1}:C{P2})", None, None], "0.0",
     "IP 授權與 EDA（Rambus、新思、益華）、手機與資料中心晶片客戶兼競爭者（高通、博通、邁威爾）；Synopsys、Cadence 已在 Bloomberg 前瞻倍數資料外，相對倍數只計 csv 內同業；非 A 股公司不列 A 股"); r_ += 1
tset(r_, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*({5/12}*{c}{EPSR[2028]}+{7/12}*{c}{EPSR[2029]})" for s, c in zip(SC, "BCD")), None, None], "0.00",
     "約 2027-10 時的未來 12 月（2027-11～2028-10）≈ 5／12 × FY2028＋7／12 × FY2029（會計年度 3 月底）"); r_ += 1
tset(r_, "comptp", "Comparable 目標價", [f"=B{T['pmed']}*B{T['feps']}", None, None], PXF, "同業中位數 × 上列 EPS"); r_ += 2
V.cell(r_, 1, "0.3 結論").font = BOLD; r_ += 1
hdr(V, r_, ["項目", "數值", "", "", "說明"]); T["h3"] = r_; r_ += 1
tset(r_, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], PXF, "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價（相對倍數調整後，見第 0.5 節；ARM 無 Bloomberg 前瞻倍數，係數＝1）"); r_ += 1
tset(r_, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], PXF, "引用 0.2（未作相對倍數調整）"); r_ += 1
tset(r_, "cons", "市場共識目標價", ["=cons_tp", None, None], PXF, "stockanalysis（S&P Global）43 位平均 290.71，2026-10-21"); r_ += 1
tset(r_, "upw", "上檔空間：加權目標價", [f"=B{T['tpw']}/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-30 收盤 228.88）"); r_ += 1
tset(r_, "upc", "上檔空間：Comparable 目標價", [f"=B{T['comp2']}/px-1", None, None], "0.0%"); r_ += 1
_u = f"B{T['upw']}"
V.cell(r_, 1, "上檔空間級距"); T["grade"] = r_
V[f"B{r_}"] = f'=IF({_u}>0.5,"大於 50%",IF({_u}>0.25,"25%–50%",IF({_u}>=0,"0%–25%",IF({_u}>=-0.25,"0%～−25%",IF({_u}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"E{r_}"] = "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）"; r_ += 1
V.cell(r_, 1, "投資建議"); T["rec"] = r_
V[f"B{r_}"] = f'=IF({_u}>0.5,"強力買進",IF({_u}>0.25,"買進",IF({_u}>=-0.25,"中立",IF({_u}>=-0.5,"放空","強力放空"))))'
V[f"E{r_}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空"; r_ += 1
tset(r_, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS（約 2027-10 起 12 個月）÷ 共識前瞻 EPS（現在起 12 個月）− 1；期間差約一年"); r_ += 1
tset(r_, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%"); r_ += 1
V[f"E{T['gmul']}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(r_, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0.00")&"～"&TEXT(B{T["dcf12"]},"#,##0.00")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）"); r_ += 1
PEROW = r_ + 3; PE10 = f"$B${PEROW + 1}:$K${PEROW + 1}"
tset(r_, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
     [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在年底 P/E 分布（下方列；2023 年 9 月上市，只有 2023–2025 三年）的分位；低於最小值記 0%、高於最大值記 100%；只作檢查，不進目標價"); r_ += 1
V.cell(r_, 1, "目標價方法"); T["m"] = r_
V[f"B{r_}"] = ("DCF：股東現金流＝non-GAAP EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年平均約 10.5%（GAAP；再投資率約 38%），FY2027–29 驅動因子、FY2030–32 收斂法＋永續；"
               "non-GAAP EPS 排除股票薪酬（FY2026 GAAP EPS 0.85 vs non-GAAP 1.77），DCF 因此偏高估，但 ROE 低使再投資率高、偏低估。Comparable：高通、博通、邁威爾、Rambus、新思、益華前瞻 P/E 中位數；"
               "Arm 現價約 95 倍共識前瞻 EPS，遠高於同業，Comparable 法等於假設溢價在 12 個月內消失，偏低估。SoftBank 持股約九成、流通股約一成，市場流通量低可能支撐溢價；本模型不作流通量調整（只在方法說明）。"
               "ARM 無 Bloomberg 前瞻倍數歷史，相對倍數係數＝1。")
r_ += 1
V.cell(PEROW, 1, "年底 P/E（歷史分位檢查用；12 月收盤 ÷ 該月所在會計年度的 non-GAAP EPS；2023 年 9 月上市，2016–2022 年無資料）").font = BOLD
CL = {2023: (75.14, 1.27), 2024: (123.36, 1.63), 2025: (109.31, 1.77)}
for j, yy in enumerate(range(2016, 2026)):
    V.cell(PEROW, 2 + j, str(yy))
    if yy in CL:
        cc = V.cell(PEROW + 1, 2 + j, f"={CL[yy][0]}/{CL[yy][1]}"); cc.number_format = "0.0"
V.cell(PEROW + 1, 1, "年底 P/E"); T["pe10h"] = PEROW; T["pe10"] = PEROW + 1
V.cell(PEROW + 2, 1, "來源：12 月收盤 digrin 月線（2023 75.14、2024 123.36、2025 109.31）；non-GAAP EPS FY2024 1.27、FY2025 1.63、FY2026 1.77（股東信）。stockanalysis GAAP P/E（會計年度末）FY2024 419.9、FY2025 142.1、FY2026 177.7 只作核對")
for k_ in ("dcftp", "comptp", "tpw"): V[f"B{T[k_]}"].fill = KEY
V.cell(T["tpw"], 1).font = BOLD
defname("tp_final", f"Val!$B${T['tpw']}"); defname("tp_dcf", f"Val!$B${T['dcftp']}"); defname("tp_comp", f"Val!$B${T['comptp']}")
v1 = PEROW + 5; T["s1"] = v1
V.cell(v1, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, v1 + 1, ["項目", "A", "B", "C", "說明"])
V.cell(v1 + 2, 1, "現價（2026-10-30 收盤）"); V[f"B{v1 + 2}"] = "=px"; V[f"B{v1 + 2}"].number_format = PXF; T["px"] = v1 + 2
for k_, (lab, y) in enumerate((("現價 ÷ FY2026 EPS（實績）", 2026), ("現價 ÷ FY2027 EPS（模型）", 2027), ("現價 ÷ FY2028 EPS（模型）", 2028), ("現價 ÷ FY2029 EPS（模型）", 2029))):
    rr = v1 + 3 + k_; V.cell(rr, 1, lab); T[f"pe{y}"] = rr
    for j, s in enumerate(SC):
        V[f"{'BCD'[j]}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{'BCD'[j]}{rr}"].number_format = "0.0"
rr = v1 + 7
V.cell(rr, 1, "近 12 月 non-GAAP EPS（2Q26～1Q27）"); V[f"B{rr}"] = "=1.77-0.35+0.45"; V[f"B{rr}"].number_format = "0.00"; T["ttm"] = rr
V.cell(rr + 1, 1, "  現價 ÷ 近 12 月 non-GAAP EPS"); V[f"B{rr + 1}"] = f"=px/B{rr}"; V[f"B{rr + 1}"].number_format = "0.0"; T["pettm"] = rr + 1
V.cell(rr + 2, 1, "  現價 ÷ 共識前瞻 EPS"); V[f"B{rr + 2}"] = "=px/cons_feps"; V[f"B{rr + 2}"].number_format = "0.0"; T["pecons"] = rr + 2
V.cell(rr + 3, 1, "  市值（現價 × 稀釋股數，十億美元）"); V[f"B{rr + 3}"] = "=px*q_sh"; V[f"B{rr + 3}"].number_format = "#,##0"; T["mcap"] = rr + 3
V.cell(rr + 4, 1, "  流通市值（× (1 − SoftBank 持股)，十億美元）"); V[f"B{rr + 4}"] = "=px*q_sh*(1-sb_own)"; V[f"B{rr + 4}"].number_format = "#,##0"; T["float"] = rr + 4
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次改一個輸入（In 頁敏感度低／高值），整本活頁簿以 LibreOffice 重算後的加權目標價（Val 第 0 節）與情境 B 的 FY2028 EPS；由 arm/sens_arm.py 產生 arm/sens.json，建置時寫入（數值）"
SENS_F = os.path.join(HERE, "sens.json")
SJ = json.load(open(SENS_F, encoding="utf-8")) if (os.path.exists(SENS_F) and not OVR) else None
LBL = {x[0]: x[1] for x in INLIST}
SE["A2"] = ("相對基準＝案例目標價 ÷ 基準加權目標價 − 1；龍捲風依絕對值排序（HTML）。FY2027 錨點與 FY2026 比例的輸入（q2_lg、ph_s 等）透過錨點影響之後各年。"
            "tp_g 案例同時移動 DCF 永續成長與 PL 收斂目標 fade_g（fade_g 以 tp_g 為公式）。")
hdr(SE, 3, ["驅動", "設定", "方向", "加權目標價", "相對基準", "鍵", "FY2028 EPS（情境 B）", "FY2028 EPS 相對基準", "說明"])
SE.cell(4, 1, "基準"); SE.cell(4, 2, "In 頁基準欄"); SE.cell(4, 3, "—"); SE.cell(4, 9, "本活頁簿重算值")
SE["D4"] = "=tp_final"; SE["G4"] = f"=PL!F{SUMR[('B', 'eps')]}"
SE["D4"].number_format = PXF; SE["G4"].number_format = "0.00"; SE["E4"] = 0; SE["H4"] = 0
for c in "EH": SE[f"{c}4"].number_format = "0.0%"
SENS_KEYS = [x[0] for x in INLIST if x[3] is not None and x[5] is not None and x[3] != x[5]]
srow = 5
for k in SENS_KEYS:
    lo, base, hi = next((x[3], x[4], x[5]) for x in INLIST if x[0] == k)
    for side, val in (("低", lo), ("高", hi)):
        SE.cell(srow, 1, f"{LBL[k]}：{side}"); SE.cell(srow, 2, f"{base} → {val}"); SE.cell(srow, 3, side); SE.cell(srow, 6, k)
        if SJ and k in SJ:
            tp_, eps_ = SJ[k][side]
            SE.cell(srow, 4, tp_).number_format = PXF; SE.cell(srow, 5, f"=D{srow}/D$4-1").number_format = "0.0%"
            SE.cell(srow, 7, eps_).number_format = "0.00"; SE.cell(srow, 8, f"=G{srow}/G$4-1").number_format = "0.0%"
            SE.cell(srow, 9, "重算值（sens_arm.py）")
        else:
            SE.cell(srow, 9, "未產生（請先執行 sens_arm.py）")
        srow += 1
SENS_LAST = srow - 1
SE.cell(srow + 1, 1, "讀法：DCF 與 Comparable 目標價都遠低於現價，DCF 權重是最大的單一變數（Comparable 約為 DCF 的兩倍）；其次是影響 FY2028–29 EPS 的授權成長、營業費用與晶片毛利率。手機與資料中心的拆分比例只改變權利金的組成，對目標價影響小。")
SENS_NOTE = srow + 1
SE.column_dimensions["A"].width = 60; SE.column_dimensions["B"].width = 18; SE.column_dimensions["I"].width = 30
for c in "CDEFGH": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司指引與共識。只作事後對照，不為對上數字而調參數。分部共識（授權、權利金分列）：未找到；FY2028 共識 EPS：未找到"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.000", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 公司指引與共識（校準未使用共識）").font = BOLD; rr_ += 1
rec("2Q27 營收 vs 公司指引中點", "=1.38", lambda s: "=Q!E6", "股東信 2026-07-29", "模型＝去年同期 × 指引年增")
rec("2Q27 營收 vs 共識", "=c_q2rev", lambda s: "=Q!E6", "Hudson Labs", "")
rec("2Q27 non-GAAP EPS vs 公司指引中點", "=0.47", lambda s: "=Q!E15", "股東信", "", fmt="0.00")
rec("FY2027 營收 vs 共識", "=c_rev27", pl("rev", 2027), "stockanalysis（S&P Global）", "", fmt="#,##0.000")
rec("FY2027 non-GAAP EPS vs 共識", "=c_eps27", pl("eps", 2027), "同上", "高 2.38、低 2.06", fmt="0.00")
RC.cell(rr_, 1, "B. 權利金結構（機制檢查）").font = BOLD; rr_ += 1
rec("FY2026 智慧型手機 ÷ 權利金（20-F）", "=ph_s", lambda s: f"=PL!D{PB[s]['phsh']}", "20-F", "", fmt="0.0%", pt=True)
rec("FY2028 資料中心 ÷ 權利金", None, lambda s: f"=PL!F{PB[s]['dcsh']}", "—", "管理層：資料中心將成為最大業務（4Q26 股東信）", fmt="0.0%")
rec("FY2029 資料中心 ÷ 權利金", None, lambda s: f"=PL!G{PB[s]['dcsh']}", "—", "", fmt="0.0%")
rec("FY2026 手機每支權利金（美元）", None, lambda s: f"=PL!D{PB[s]['phrt']}", "—", "Derived＝43% × 權利金 ÷ 出貨", fmt="0.000")
rec("FY2028 NVIDIA Grace／Vera CPU（百萬顆）", None, lambda s: f"=PL!F{PB[s]['nvcpu']}", SNAP, "", fmt="0.00")
rec("FY2028 AGI CPU 晶片營收（US$B）", None, pl("chip", 2028), "—", "客戶需求 FY2027–28 逾 20 億美元（公司）", fmt="0.000")
rec("FY2027 + FY2028 AGI CPU 晶片營收合計 vs 公司稱客戶需求", "=2", lambda s: f"=PL!E{PB[s]['chip']}+PL!F{PB[s]['chip']}", "股東信 2026-07-29", "需求≠營收（產能限制）", fmt="0.000")
RECON_LAST = rr_ - 1

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝Arm 股東信、6-K、20-F、法說；共識與券商皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：1Q27 股東信（6-K Exhibit 99.2，2026-07-29；1Q27 實績、2Q27 指引、AGI CPU 需求、資料中心權利金）", SL27),
        ("公司：4Q26 股東信（6-K，2026-05-06；FY2026 全年、AGI CPU、1Q27 指引）", SL26),
        ("公司：4Q25 股東信（6-K，2025-05；FY2025 全年、FY2024 non-GAAP EPS）", SL25),
        ("公司：20-F（FY2026；手機權利金約 43%、前五大客戶 57%、Arm China 16%、SoftBank 控股）", F20),
        ("公司：3Q26 法說逐字稿（The Motley Fool 2026-02-04；資料中心占比、CSS、SoftBank 協議 2 億美元／季、FY2027 約 20%）", TR3),
        ("公司：Arm 新聞室 1Q27 結果（2026-07-29）", "https://newsroom.arm.com/news/arm-q1-fye27-results"),
        ("公司：2Q27 財報日公告（2026-11-04 美東盤後）", "https://newsroom.arm.com/news/arm-announces-earnings-release-date-for-second-quarter-fiscal-year-ending-2027"),
        ("媒體：Longbridge 2026-07-29（2Q27 指引與共識）", LB),
        ("研究：Hudson Labs 2Q27 財報預覽（共識、FY2027 手機權利金 high-teens）", HL),
        ("法律：Paul Weiss 2026-10-13（Qualcomm 對 Arm 陪審團判決）", PW),
        ("資料商：stockanalysis（共識 EPS、目標價、Forward PE、ROE；同業 Forward PE）", SAF),
        ("資料商：InvestSmart 價格歷史（2026-10-30 收盤 228.88）；digrin 月線（年底收盤）", IS),
        ("市場層：Siliconomics 快照（NVIDIA 加速器顆數、情境顆數係數）", SNAP),
        ("待補（人工）：2Q27 財報（2026-11-04）；分市場權利金（公司只揭露手機）；AGI CPU 營收認列方式；Qualcomm 法官審理判決；FY2028 共識；Bloomberg 前瞻倍數（ARM 未取）", "未找到")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
OBS = [   # 重點觀察（數字取自重算後的 Val／PL／Sens；與 README.md 相同文字）
 "目標價 US$53.52（DCF 34.56、Comparable 72.48 各半），比現價 228.88（2026-10-30）低 76.6%，級距小於 −50%、投資建議強力放空；Comparable 單獨為 −68.3%。"
 "差距全部來自倍數：本模型 12 個月後前瞻 EPS 2.90 反而比共識前瞻 EPS 高 20.1%（期間晚約一年），但現價約 94.7 倍共識前瞻 EPS、共識目標價 290.71 隱含 120.3 倍，同業前瞻 P/E 中位數只有 25.0 倍。"
 "DCF 偏低另因 GAAP ROE 近三年平均 10.5%（再投資率 38%）。SoftBank 持股約九成，流通市值約 247 億美元（總市值 2,472 億），流通量低可能撐住溢價，本模型不調整；ARM 無 Bloomberg 前瞻倍數，相對倍數係數＝1。",
 "單一參數都翻不過來：最有利的單一案例是 DCF 權重 0.3（+14.2%），其次授權成長 15%（+6.8%）、營業費用年增 8%（+7.4%），目標價仍比現價低七成以上；授權成長 2% −10.3%、營業費用年增 18% −7.7%、晶片毛利率 35%／55% ±4.6%（Sens 頁）。",
 "可追蹤的三件事：資料中心占權利金由 FY2026 14% 升到 FY2029 情境 B 24.9%（Grace／Vera 讀 Siliconomics、雲端自研 CPU 年增 30%）；AGI CPU 晶片營收 FY2028 情境 A／B／C＝14／10／5 億美元，對照公司稱 FY2027–28 客戶需求逾 20 億美元；Qualcomm 陪審團 2026-10-13 判 Arm 違約，法官審理部分待判。下一個檢驗點：2Q27 財報 2026-11-04（指引營收 13.8 億美元、non-GAAP EPS 0.47）。"]
lines = [(f"Arm Holdings（ARM）個股模型 {VER}（AI 半導體 Project：運算次產業；CPU 架構授權與權利金）　{DATE}", True),
 (f"版本：{VER}（{DATE}，初版）。樣板 crdo/（美股、估值 v2）；敏感度與份額寫法比照 tuc/；行業變數讀 {SNAP}。待 FQ2 財報（2026-11-04）更新錨點。", False),
 ("模型命題：Arm 的權利金＝各市場晶片出貨 × 每顆權利金率：智慧型手機（出貨持平到下滑，成長靠 v9 與 CSS 費率）、資料中心 Neoverse（NVIDIA Grace／Vera 顆數讀 Siliconomics、雲端自研 CPU 年增、每顆費率隨核心數上升）、車用、物聯網與其他；授權及其他＝授權（含 SoftBank 協議）＋自研 Arm AGI CPU 晶片。", False),
 ("核心問題：現價約 95 倍共識前瞻 EPS，市場是否已把資料中心權利金倍增、CSS 費率上升與自研晶片（客戶需求逾 20 億美元）全部算進去；Qualcomm 判決與手機出貨下修會讓權利金少多少。", False),
 ("會計年度：FY(t) 於 t 年 3 月 31 日結束（FY2027＝2026-04～2027-03）；上游日曆年以 9/12 × CY(t−1)＋3/12 × CY(t) 換算。預測期 FY2028–FY2029 由驅動因子推導，FY2030–FY2032 為收斂法。金額：十億美元；EPS：non-GAAP（GAAP EPS 另列）。", False),
 ("驅動因子：FY2027 錨點＝1Q27 實績＋2Q27 指引＋下半年年增假設；權利金分市場以 FY2026 比例（手機 43% 為 20-F，其餘 Assumed）拆分；FY2028 起資料中心＝FY2027 × CPU 顆數指數（NVIDIA 讀 Siliconomics 同字母情境）× 每顆費率指數；手機＝出貨 × 每支權利金（情境 A／B／C 費率年增 14%／11%／7%）；晶片營收情境 A／B／C。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），FY2027–FY2032 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，推到 12 個月後）＋0.5 × Comparable 目標價（高通、博通、邁威爾、Rambus、新思、益華前瞻 P/E 中位數 × 12 個月後前瞻 EPS）。本版未做先驗，無 Andy 合理 P/E 對照。Bloomberg 前瞻倍數待取（ARM 本身沒有，相對倍數係數＝1）。", False),
 ("已知限制：(1) 公司只揭露手機占權利金比例，其他市場拆分為 Assumed；(2) AGI CPU 晶片營收所在列與毛利率未揭露；(3) non-GAAP EPS 排除股票薪酬（FY2026 約占 non-GAAP 淨利一半）；(4) Qualcomm 法官審理部分未判、損害賠償未知；(5) 年底 P/E 只有 2023–2025 三年；(6) 2Q27 財報（2026-11-04）未納入。", False),
 ("", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
RMX.apply(wb, "arm")
apply_in_layout.apply(wb, VER, DATE)
for row_ in RD.iter_rows():
    for c_ in row_:
        if isinstance(c_.value, str) and c_.value.startswith("2026-10-29：加權目標價中的 Comparable"):
            c_.value = "相對倍數調整：加權目標價中的 Comparable 用相對倍數調整後的值（Val 第 0.5 節；tools/relative_multiple.py）。ARM 無 Bloomberg 前瞻倍數歷史（待取），係數＝1（不調整）。「Comparable 目標價（市場主流同業法）」不調整。"
        elif isinstance(c_.value, str) and c_.value.startswith(f"{VER}（{DATE}）：依個股模型原則第 8 節"):
            c_.value = "In 頁版面（個股模型原則第 8 節）：D／F 欄＝敏感度低／高，只用於 Sens 頁、不進目標價；In 頁最後一節「情境驅動對照」把情境專屬輸入三情境並排（只引用上方各列，純展示）。"
rn = RD.max_row + 2; OBSROW = rn
RD.cell(rn, 1, f'="投資建議："&Val!B{T["rec"]}&"（依上檔空間固定規則；加權目標價上檔空間 "&TEXT(Val!B{T["upw"]},"+0.0%;−0.0%")&"）"'); rn += 1
for k_, o_ in enumerate(OBS, 1):
    RD.cell(rn, 1, f"重點觀察 {k_}：" + o_); rn += 1
wb.save(sys.argv[1])
if os.environ.get("ARM_ROWMAP"):
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "XT": XT, "XH": XH, "XR": XR, "T": T, "EPSR": EPSR, "CFR": CFR,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "UNOTE": UNOTE, "INROW": INROW, "OBSROW": OBSROW, "SENS_KEYS": SENS_KEYS, "QR": QR},
              open(os.environ["ARM_ROWMAP"], "w"), ensure_ascii=False, indent=1)
