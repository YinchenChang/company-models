# Arm Sens 產生器：一次把 In 頁一個鍵的基準值換成敏感度低／高值，整本重建、LibreOffice 重算，讀加權目標價（tp_final）與情境 B FY2028 EPS，
# 寫入 arm/sens.json（建置時 build_arm.py 讀入 Sens 頁）。In 頁或模型改動後重跑一次，再正式建置。
# 用法：python3 arm/sens_arm.py [暫存資料夾，預設 /tmp/work_arm/sens]
import json, os, subprocess, sys, glob, shutil
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
TMP = sys.argv[1] if len(sys.argv) > 1 else "/tmp/work_arm/sens"
shutil.rmtree(TMP, ignore_errors=True); os.makedirs(os.path.join(TMP, "raw")); os.makedirs(os.path.join(TMP, "rc"))
BUILD = os.path.join(HERE, "build_arm.py")
env0 = {k: v for k, v in os.environ.items() if k not in ("ARM_OVERRIDE", "ARM_ROWMAP")}
base_raw = os.path.join(TMP, "raw", "base.xlsx")
subprocess.run([sys.executable, BUILD, base_raw], check=True, env={**env0, "ARM_OVERRIDE": "_none=0"})
ws = openpyxl.load_workbook(base_raw)["In"]
jobs = {}
for r in range(4, ws.max_row + 1):
    k, lo, hi = ws.cell(r, 1).value, ws.cell(r, 4).value, ws.cell(r, 6).value
    if k and isinstance(lo, (int, float)) and isinstance(hi, (int, float)) and lo != hi:
        if "情境驅動對照" in str(ws.cell(r - 1, 1).value or ""): break
        jobs[k] = (lo, hi)
for k, (lo, hi) in jobs.items():
    for side, v in (("lo", lo), ("hi", hi)):
        subprocess.run([sys.executable, BUILD, os.path.join(TMP, "raw", f"{k}__{side}.xlsx")], check=True, env={**env0, "ARM_OVERRIDE": f"{k}={v}"})
files = sorted(glob.glob(os.path.join(TMP, "raw", "*.xlsx")))
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(TMP, "rc")] + files,
               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
def read(f):
    wb = openpyxl.load_workbook(f, data_only=True)
    dn = wb.defined_names["tp_final"].attr_text.split("!")[1].replace("$", "")
    tp = wb["Val"][dn].value
    eps = wb["Sens"]["G4"].value
    return tp, eps
out = {}
b = read(os.path.join(TMP, "rc", "base.xlsx"))
for k in jobs:
    out[k] = {"低": read(os.path.join(TMP, "rc", f"{k}__lo.xlsx")), "高": read(os.path.join(TMP, "rc", f"{k}__hi.xlsx"))}
json.dump(out, open(os.path.join(HERE, "sens.json"), "w"), ensure_ascii=False, indent=1)
print("基準", b, "；案例", len(out))
for k, v in sorted(out.items(), key=lambda kv: -max(abs(kv[1]["低"][0] / b[0] - 1), abs(kv[1]["高"][0] / b[0] - 1))):
    print(f"{k:10s} 低 {v['低'][0] / b[0] - 1:+.1%}  高 {v['高'][0] / b[0] - 1:+.1%}")
