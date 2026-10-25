# 目標價 v2（Andy 2026-10-25 11:28、11:53）：DCF（股東權益現金流）與 Comparable（同業前瞻 P/E）加權；三情境延伸到 FY2032（收斂法）。
# 本檔在各公司資料夾各放一份（內容相同），由 build_<公司>.py 在存檔前呼叫；不讀寫其他資料夾。
#   extend_pl(...)      PL 頁末加「三情境延伸到 FY2032」表（上游結束後用收斂法）
#   shift_rows(...)     把 Val 頁既有內容整體下移，公式、定義名稱跨頁一併改寫（讓第 0 節放在 Val 最前面）
#   add_tp_section(...) Val 頁最前面寫「0. 目標價」（0.1 DCF、0.2 Comparable、0.3 結論；列標籤依工作單逐字）
#   add_tp_sens(...)    Sens 頁加 tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.3／0.7 對加權目標價的影響
import re
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True); BLUE = Font(color="0000FF"); ITAL = Font(italic=True)
HDR = PatternFill("solid", fgColor="DDEBF7"); KEY = PatternFill("solid", fgColor="FFF2CC")
CONV = PatternFill("solid", fgColor="FCE4D6")   # 收斂法年度底色
TP_KEYS = ("tp_wA", "tp_wB", "tp_wC", "tp_r", "tp_g", "tp_roe", "tp_w_dcf", "tp_gd", "nm10", "nm10_k")
SC = ("A", "B", "C")


def hdr(ws, row, labels, c0=1):
    for j, h in enumerate(labels, c0):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR


def yellow_inputs(IN, keys=TP_KEYS):
    """In 頁：目標價參數列塗黃底（Andy 可改）。"""
    for r in range(4, IN.max_row + 1):
        if IN.cell(r, 1).value in keys:
            for c in range(1, 10): IN.cell(r, c).fill = KEY


# ---------------- 公式參照位移 ----------------
_REF = re.compile(r"(?<![A-Za-z0-9_.$!])(?:(?P<sh>'[^']+'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
                  r"(?P<c1>\$?[A-Z]{1,3})(?P<d1>\$?)(?P<n1>\d+)(?::(?P<c2>\$?[A-Z]{1,3})(?P<d2>\$?)(?P<n2>\d+))?(?![A-Za-z0-9_(!])")


def _fix(formula, cur, sheet, start, n):
    parts = formula.split('"')
    for i in range(0, len(parts), 2):          # 只改字串常值以外的部分
        def rep(m):
            sh = m.group("sh")
            tgt = sh.strip("'") if sh else cur
            if tgt != sheet:
                return m.group(0)
            n1 = int(m.group("n1")); n1 = n1 + n if n1 >= start else n1
            s = (sh + "!" if sh else "") + m.group("c1") + m.group("d1") + str(n1)
            if m.group("n2"):
                n2 = int(m.group("n2")); n2 = n2 + n if n2 >= start else n2
                s += ":" + m.group("c2") + m.group("d2") + str(n2)
            return s
        parts[i] = _REF.sub(rep, parts[i])
    return '"'.join(parts)


def shift_rows(wb, sheet, start, n):
    """sheet 第 start 列起整體下移 n 列；全活頁簿公式與定義名稱中指向該區的參照一併改寫。"""
    ws = wb[sheet]
    assert not ws.merged_cells.ranges and ws.freeze_panes in (None, "A1"), (sheet, ws.freeze_panes)
    ws.move_range(f"A{start}:{get_column_letter(ws.max_column)}{ws.max_row}", rows=n, translate=False)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = _fix(c.value, w.title, sheet, start, n)
    for nm, dn in list(wb.defined_names.items()):
        dn.attr_text = _fix(dn.attr_text, None, sheet, start, n)


# ---------------- PL：三情境延伸到 FY2032 ----------------
def extend_pl(PL, blocks, cols, years, last_up, keys, unit_px, title, conv_note, m10="nm10_k"):
    """blocks：{s: 引擎列號 dict}；cols：{FY: 欄}（引擎的年度欄）；years：引擎年度（含實績年）；last_up：上游最後一年（之後為收斂法）；
    keys：{"rev","ni","sh","eps"} 對應引擎鍵。回傳 (表頭列, {(s, 指標): 列}, {FY: 欄})。"""
    r0 = PL.max_row + 3
    allys = list(years) + list(range(last_up + 1, 2033))
    XC = {y: get_column_letter(3 + i) for i, y in enumerate(allys)}
    for y in years: assert XC[y] == cols[y], (y, XC[y], cols[y])     # 延伸表的年度欄與引擎同欄，FY2032 之前逐欄對齊
    note_c = 3 + len(allys)
    PL.cell(r0, 1, title).font = BOLD
    PL.cell(r0 + 1, 1, conv_note)
    hr = r0 + 2
    hdr(PL, hr, ["情境／延伸項目", "單位"] + [f"FY{y}" + ("A" if y == years[0] else "") for y in allys] + ["說明"])
    conv = [y for y in allys if y > last_up]; N = len(conv)
    R = {}; r = hr + 1
    for s in SC:
        B = blocks[s]
        PL.cell(r, 1, f"情境 {s}").font = BOLD; r += 1
        rr = {k: r + i for i, k in enumerate(("rev", "g", "nm", "ni", "sh", "eps", "epsg"))}
        R.update({(s, k): v for k, v in rr.items()})
        labs = {"rev": ("營收", PL.cell(B[keys["rev"]], 2).value), "g": ("  營收年增", "%"), "nm": ("淨利率", "%"),
                "ni": ("淨利（延伸表）", PL.cell(B[keys["ni"]], 2).value), "sh": ("稀釋股數", PL.cell(B[keys["sh"]], 2).value),
                "eps": ("EPS（延伸表）", unit_px), "epsg": ("  EPS 年增", "%")}
        for k, (lab, u) in labs.items():
            PL.cell(rr[k], 1, lab); PL.cell(rr[k], 2, u)
        lc = XC[last_up]
        gT = {"A": "(tp_g+tp_gd)", "B": "tp_g", "C": "(tp_g-tp_gd)"}[s]
        M10 = "nm10" if m10 == "nm10" else f"${lc}${rr['nm']}*nm10_k"     # 十年平均淨利率；找不到同口徑者用最後一年 × nm10_k（0.8，Assumed）
        mT = {"A": f"${lc}${rr['nm']}", "B": f"(${lc}${rr['nm']}+{M10})/2", "C": M10}[s]
        for i, y in enumerate(allys):
            c = XC[y]; p = XC.get(y - 1)
            if y <= last_up:
                for k in ("rev", "ni", "sh", "eps"):
                    PL[f"{c}{rr[k]}"] = f"={c}{B[keys[k]]}"
                if p: PL[f"{c}{rr['g']}"] = f"={c}{rr['rev']}/{p}{rr['rev']}-1"
                PL[f"{c}{rr['nm']}"] = f"={c}{rr['ni']}/{c}{rr['rev']}"
            else:
                k_ = y - last_up
                PL[f"{c}{rr['g']}"] = f"=${lc}${rr['g']}+({gT}-${lc}${rr['g']})*{k_}/{N}"
                PL[f"{c}{rr['rev']}"] = f"={p}{rr['rev']}*(1+{c}{rr['g']})"
                PL[f"{c}{rr['nm']}"] = f"=${lc}${rr['nm']}+({mT}-${lc}${rr['nm']})*{k_}/{N}"
                PL[f"{c}{rr['ni']}"] = f"={c}{rr['rev']}*{c}{rr['nm']}"
                PL[f"{c}{rr['sh']}"] = f"={p}{rr['sh']}*(1+${lc}${rr['sh']}/${XC[last_up - 1]}${rr['sh']}-1)"
                PL[f"{c}{rr['eps']}"] = f"={c}{rr['ni']}/{c}{rr['sh']}"
                for k in rr: PL[f"{c}{rr[k]}"].fill = CONV
            if p: PL[f"{c}{rr['epsg']}"] = f"={c}{rr['eps']}/{p}{rr['eps']}-1"
            for k, f in (("rev", "#,##0.0"), ("g", "0.0%"), ("nm", "0.0%"), ("ni", "#,##0.0"), ("sh", "0.000"), ("eps", "#,##0.00"), ("epsg", "0.0%")):
                PL[f"{c}{rr[k]}"].number_format = f
        PL.cell(rr["g"], note_c, f"FY{last_up + 1} 起：上游最後一年（FY{last_up}）年增線性收斂到 FY2032 的長期成長（" + {"A": "tp_g＋tp_gd", "B": "tp_g", "C": "tp_g－tp_gd"}[s] + "）")
        PL.cell(rr["nm"], note_c, f"FY{last_up + 1} 起：FY{last_up} 淨利率線性收斂到中段利潤率（" + {"A": "維持最後一年", "B": "最後一年與十年平均的中點", "C": "十年平均"}[s] + "；十年平均見 In 頁 " + ("nm10" if m10 == "nm10" else "nm10_k") + " 列）")
        PL.cell(rr["sh"], note_c, f"FY{last_up + 1} 起沿用 FY{last_up} 的股數年變動")
        r = rr["epsg"] + 2
    return hr, R, XC


# ---------------- Val 第 0 節 ----------------
def add_tp_section(wb, V, eps, e_mon, fwd, fwd_note, peers, peer_note, cons_tp, cons_note, cons_fwd, cons_fwd_note,
                   eps28, hist, fmt, method, start=3, dcf_note="", ours=None):
    """eps：{s: [FY2027..FY2032 的 PL 儲存格]}；e_mon：FY2027 年底距 2026-12-31 的月數；fwd：[(PL 欄位 dict 的鍵, 權重)] → {s: [(cell, w)]}；
    peers：[(名稱, 代號, 前瞻 P/E, 日期, 來源, 標記)]；cons_tp：In 名稱或 None；cons_fwd：共識前瞻 EPS 公式或 None；
    eps28：{s: PL 儲存格}；hist：位移前的 Val 範圍（如 'B80:B89'）。回傳 (列號 dict, 位移列數)。"""
    npeer = len(peers)
    N = 46 + npeer
    shift_rows(wb, "Val", start, N)
    hm = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", hist)
    hist_s = f"{hm.group(1)}{int(hm.group(2)) + N}:{hm.group(3)}{int(hm.group(4)) + N}"
    r = start; L = {}
    def put(lab, vals=None, note=None, key=None, f=None, bold=False, fill=None):
        nonlocal r
        V.cell(r, 1, lab)
        if bold: V.cell(r, 1).font = BOLD
        for j, v in enumerate(vals or []):
            if v is None: continue
            c = V.cell(r, 2 + j, v)
            if f: c.number_format = f[j] if isinstance(f, (list, tuple)) else f
        if note is not None: V.cell(r, 5, note)
        if fill:
            for cc in range(1, 6): V.cell(r, cc).fill = fill
        if key: L[key] = r
        r += 1
    put("0. 目標價（DCF 與 Comparable 加權；v0.2 新增，Andy 2026-10-25）", bold=True)
    put("0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）", bold=True)
    hdr(V, r, ["項目", "A", "B", "C", "說明", "折現年數 t", "期內占比（基準日後）"]); L["h1"] = r; r += 1
    put("情境 EPS（FY2027–FY2032）", note="引用 PL 頁延伸表（上游結束後為收斂法）；三情境欄 B／C／D＝A／B／C", key="eps_h")
    ts = [(e_mon + 12 * k) / 12 for k in range(6)]; fr = [e_mon / 12] + [1] * 5
    for k in range(6):
        put(f"  FY{2027 + k}", [f"={eps[s][k]}" for s in SC], f="#,##0.00", key=f"eps{k}")
        V.cell(r - 1, 6, round(ts[k], 6)).number_format = "0.000"; V.cell(r - 1, 7, round(fr[k], 6)).number_format = "0.000"
    put("股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）", ["=1-MIN(0.8,tp_g/tp_roe)"], f="0.0%", key="po",
        note="B 欄＝可發放比例（1 − 再投資率；再投資率＝tp_g ÷ tp_roe，上限 0.8）；下列六列＝EPS × 可發放比例 × 期內占比" + dcf_note)
    for k in range(6):
        put(f"  FY{2027 + k} 股東現金流", [f"={c}{L[f'eps{k}']}*$B${L['po']}*$G${L[f'eps{k}']}" for c in "BCD"], f="#,##0.00", key=f"cf{k}")
    put("  永續價值（FY2032 年底；FY2032 現金流 ×(1＋tp_g) ÷ (tp_r − tp_g)）", [f"={c}{L['cf5']}*(1+tp_g)/(tp_r-tp_g)" for c in "BCD"], f="#,##0.00", key="tv")
    pv = lambda c: "+".join(f"{c}{L[f'cf{k}']}/(1+tp_r)^$F${L[f'eps{k}']}" for k in range(6)) + f"+{c}{L['tv']}/(1+tp_r)^$F${L['eps5']}"
    put("DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,{pv(c)})" for c in "BCD"], f=fmt, key="dcfv",
        note="基準日 2026 年底；明確期現金流與永續價值依 F 欄年數折現；小於 0 時取 0（Andy 2026-10-25 12:22）", fill=KEY)
    put("  其中：永續價值現值占比", [f"={c}{L['tv']}/(1+tp_r)^$F${L['eps5']}/{c}{L['dcfv']}" for c in "BCD"], f="0.0%", key="tvsh")
    put("DCF 目標價（情境加權）", [f"=(tp_wA*B{L['dcfv']}+tp_wB*C{L['dcfv']}+tp_wC*D{L['dcfv']})*(1+tp_r)"], f=fmt, key="dcf",
        note="Σ 權重 × 情境價值 ×(1＋tp_r)：推到約 12 個月後", fill=KEY)
    r += 1
    put("0.2 Comparable（市場主流同業前瞻 P/E）", bold=True)
    hdr(V, r, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); L["ph"] = r; r += 1
    p0 = r
    for nm, tk, pe, dt, src, tag in peers:
        V.cell(r, 1, nm); V.cell(r, 2, tk); c = V.cell(r, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(r, 4, dt); V.cell(r, 5, src); V.cell(r, 6, tag); r += 1
    L["p0"], L["p1"] = p0, r - 1
    put("  同業選擇", note=peer_note, key="pnote")
    put("同業前瞻 P/E 中位數", [f"=MEDIAN(C{L['p0']}:C{L['p1']})"], f="0.0", key="med", fill=KEY)
    put("  各情境 12 個月後的前瞻 EPS", ["=" + "+".join(f"{w}*{cell}" for cell, w in fwd[s]) for s in SC], f="#,##0.00", key="fwds", note=fwd_note)
    put("12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*B{L['fwds']}+tp_wB*C{L['fwds']}+tp_wC*D{L['fwds']}"], f="#,##0.00", key="fwd")
    put("Comparable 目標價", [f"=B{L['med']}*B{L['fwd']}"], f=fmt, key="comp", note="同業前瞻 P/E 中位數 × 12 個月後的前瞻 EPS（情境加權）", fill=KEY)
    r += 1
    put("0.3 結論", bold=True)
    hdr(V, r, ["項目", "數值", "對照", "對照", "說明"]); L["h3"] = r; r += 1
    put("目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{L['dcf']}+(1-tp_w_dcf)*B{L['comp']}"], f=fmt, key="tp", fill=KEY,
        note='="DCF 權重 "&TEXT(tp_w_dcf,"0%")&"、Comparable 權重 "&TEXT(1-tp_w_dcf,"0%")&"（In 頁 tp_w_dcf，Andy 可改）"')
    put("Comparable 目標價（市場主流同業法）", [f"=B{L['comp']}"], f=fmt, key="tpc")
    put("市場共識目標價", [f"={cons_tp}" if cons_tp else "未找到"], f=fmt, key="cons", note=cons_note)
    put("上檔空間：加權目標價", [f"=B{L['tp']}/px-1"], f="0.0%", key="up")
    put("上檔空間：Comparable 目標價", [f"=B{L['tpc']}/px-1"], f="0.0%", key="upc")
    if cons_tp and cons_fwd:
        mine = ours(L) if ours else f"B{L['fwd']}"
        put("與市場共識目標價的差距：EPS 因素", [f"={mine}/C{r}-1", f"={cons_fwd}"], f=["0.0%", "#,##0.00"], key="geps",
            note="B＝我們的前瞻 EPS（情境加權）÷ 共識前瞻 EPS − 1；C＝共識前瞻 EPS（" + cons_fwd_note + "）")
        put("與市場共識目標價的差距：倍數因素", [f"=C{r}/B{L['med']}-1", f"=B{L['cons']}/C{L['geps']}", f"=B{L['med']}"], f=["0.0%", "0.0", "0.0"], key="gmul",
            note="C＝市場共識目標價 ÷ 共識前瞻 EPS 隱含的倍數；D＝同業中位數；B＝C ÷ D − 1")
    else:
        put("與市場共識目標價的差距：EPS 因素", ["未找到"], key="geps", note="共識目標價或共識前瞻 EPS 未找到")
        put("與市場共識目標價的差距：倍數因素", ["未找到"], key="gmul", note="同上")
    put("價格帶（DCF 不如預期～順利）", [f"=D{L['dcfv']}*(1+tp_r)", f"=B{L['dcfv']}*(1+tp_r)"], f=fmt, key="band",
        note="B＝C 情境 DCF 價值、C＝A 情境 DCF 價值（皆推到 12 個月後）")
    pr = lambda x: f"IF({x}<MIN({hist_s}),0,IF({x}>MAX({hist_s}),1,PERCENTRANK({hist_s},{x})))"
    put("  檢查用：現價 ÷ 各情境 FY2028 EPS", [f"=px/{eps28[s]}" for s in SC], f="0.0", key="pe28")
    put("歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"={pr(c + str(r - 1))}" for c in "BCD"], f="0%", key="pct",
        note=f"上一列在 2016–2025 年底 P/E 分布（Val!{hist_s}）的分位（PERCENTRANK；低於最小＝0、高於最大＝1）；只作檢查，不進目標價")
    put("目標價方法", [method], key="meth")
    put("  檢查：情境權重合計（應為 1）", ["=ROUND(tp_wA+tp_wB+tp_wC,6)"], f="0.000", key="wsum")
    assert r <= start + N - 1, (r, start, N)
    return L, N


# ---------------- Sens：目標價參數 ----------------
def add_tp_sens(SE, row, L, fmt):
    """Sens 頁 row 起五列：加權目標價對 tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.3／0.7 的影響（其餘為 In 頁基準；PL 收斂路徑不隨 tp_g 變動）。"""
    eps = lambda c, k: f"Val!{c}{L[f'eps{k}']}"
    t = lambda k: f"Val!$F${L[f'eps{k}']}"; fr = lambda k: f"Val!$G${L[f'eps{k}']}"
    def V(c, rr, gg):
        po = f"(1-MIN(0.8,{gg}/tp_roe))"
        s = "+".join(f"{eps(c, k)}*{fr(k)}*{po}/(1+{rr})^{t(k)}" for k in range(6))
        return "MAX(0," + s + f"+{eps(c, 5)}*{po}*(1+{gg})/({rr}-{gg})/(1+{rr})^{t(5)})"
    dcf = lambda rr, gg: f"(tp_wA*({V('B', rr, gg)})+tp_wB*({V('C', rr, gg)})+tp_wC*({V('D', rr, gg)}))*(1+{rr})"
    tpf = lambda rr, gg, w: f"={w}*{dcf(rr, gg)}+(1-{w})*Val!$B${L['comp']}"
    SE.cell(row, 1, "目標價參數敏感度（v0.2）：加權目標價（Val 第 0 節）一次動一個參數；PL 延伸期的收斂路徑不隨 tp_g 變動，tp_g 只改再投資率與永續成長").font = BOLD
    hdr(SE, row + 1, ["參數", "設定（低／高）", "低：加權目標價", "低：相對基準", "高：加權目標價", "高：相對基準", "基準加權目標價"])
    cases = [("tp_r ±1pt", "折現率 −1pt／＋1pt", ("(tp_r-0.01)", "tp_g", "tp_w_dcf"), ("(tp_r+0.01)", "tp_g", "tp_w_dcf")),
             ("tp_g ±1pt", "長期成長 −1pt／＋1pt", ("tp_r", "(tp_g-0.01)", "tp_w_dcf"), ("tp_r", "(tp_g+0.01)", "tp_w_dcf")),
             ("tp_w_dcf 0.3／0.7", "DCF 權重 0.3／0.7", ("tp_r", "tp_g", "0.3"), ("tp_r", "tp_g", "0.7"))]
    for i, (nm, st, lo, hi) in enumerate(cases):
        rr = row + 2 + i
        SE.cell(rr, 1, nm); SE.cell(rr, 2, st)
        SE.cell(rr, 3, tpf(*lo)).number_format = fmt; SE.cell(rr, 5, tpf(*hi)).number_format = fmt
        SE.cell(rr, 7, f"=Val!$B${L['tp']}").number_format = fmt
        SE.cell(rr, 4, f"=C{rr}/G{rr}-1").number_format = "0.0%"; SE.cell(rr, 6, f"=E{rr}/G{rr}-1").number_format = "0.0%"
    return row + 4


def dump_rowmap(L, N, XH, XR, XC, sens_row):
    """除錯用：環境變數 TP_ROWMAP 設定時輸出列號對照（供撰寫 html.json）；CI 不設此變數。"""
    import os, json
    if os.environ.get("TP_ROWMAP"):
        json.dump({"L": L, "N": N, "XH": XH, "XR": {f"{s}|{k}": v for (s, k), v in XR.items()}, "XC": {str(k): v for k, v in XC.items()},
                   "SENS": sens_row}, open(os.environ["TP_ROWMAP"], "w"))
