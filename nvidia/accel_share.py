# 加速器金額份額（可重用）：同一份顆數（markets/memory 的 UNITS）× 各平台每顆售價 → 各家的金額與份額。
# NVIDIA 模型（Units 頁）用它；之後 AVGO、AMD 模型可複製使用（本 PR 只建在 nvidia/ 內）。
# 純函式：只寫公式到指定工作表，回傳列號對照；顆數與價格以「儲存格公式字串」傳入，不讀檔案。
from openpyxl.styles import Font, PatternFill

BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")

# 廠商分組：NVIDIA 四個世代、AMD、TPU、Trainium、其他 ASIC
VENDORS = [("nv", "NVIDIA", ("bw", "rb", "ru", "hp")), ("amd", "AMD", ("amd",)), ("tpu", "Google TPU", ("tpu",)),
           ("trn", "AWS Trainium", ("trn",)), ("oth", "其他 ASIC", ("oth",))]


def write_share_block(ws, r0, title, years, ycols, units, prices, platforms, pnames):
    """在 ws 的 r0 列起寫一個情境的金額份額區塊。
    years：[2026, …]；ycols：{年: 欄字母}；units：{(平台, 年): 顆數公式字串（不含等號）}；
    prices：{平台: 每顆售價公式字串（US$K，不含等號）}；platforms：平台鍵順序；pnames：{平台: 名稱}。
    回傳 (下一個空列, 列號對照 R)；R[('u'|'d', 平台)]＝該平台顆數／金額列，R[('us'|'ds', 廠商)]＝顆數／金額份額列，
    R['utot']、R['dtot']＝合計列。金額＝百萬顆 × US$K ＝ US$B。"""
    ws.cell(r0, 1, title).font = BOLD
    R = {}
    r = r0 + 1
    for kind, lab, unit in (("u", "顆數", "百萬顆"), ("d", "金額（顆數 × 每顆售價）", "US$B")):
        ws.cell(r, 1, lab).font = BOLD
        r += 1
        first = r
        for p in platforms:
            ws.cell(r, 1, f"  {pnames[p]}")
            ws.cell(r, 2, unit)
            for y in years:
                f = units[(p, y)] if kind == "u" else f"{units[(p, y)]}*({prices[p]})"
                ws[f"{ycols[y]}{r}"] = "=" + f
                ws[f"{ycols[y]}{r}"].number_format = "0.00" if kind == "u" else "0.0"
            R[(kind, p)] = r
            r += 1
        ws.cell(r, 1, "  合計").font = BOLD
        ws.cell(r, 2, unit)
        for y in years:
            c = ycols[y]
            ws[f"{c}{r}"] = f"=SUM({c}{first}:{c}{r-1})"
            ws[f"{c}{r}"].number_format = "0.00" if kind == "u" else "0.0"
        R[kind + "tot"] = r
        r += 1
    for kind, lab in (("u", "顆數份額"), ("d", "金額份額")):
        ws.cell(r, 1, lab).font = BOLD
        r += 1
        for vk, vn, members in VENDORS:
            ws.cell(r, 1, f"  {vn}")
            ws.cell(r, 2, "%")
            for y in years:
                c = ycols[y]
                num = "+".join(f"{c}{R[(kind, p)]}" for p in members if (kind, p) in R)
                ws[f"{c}{r}"] = f"=({num})/{c}{R[kind + 'tot']}"
                ws[f"{c}{r}"].number_format = "0.0%"
            R[(kind + "s", vk)] = r
            r += 1
    return r + 1, R
