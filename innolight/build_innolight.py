# 中際旭創（300308.SZ／3308.HK）收支與估值模型 v0.2.2（AI 半導體 Project：網通次產業內複製——光模組）
# 用法：python3 innolight/build_innolight.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不做平移。金額：十億人民幣（RMB B）；市場：十億美元（US$B，讀 markets/network）；EPS：人民幣。
# 口徑：中國會計準則；淨利＝歸屬母公司股東淨利。公司不給營收或獲利指引；FY2026 錨點＝1H26 實績＋2H26 假設（Q3 季報 10-31 前公布後更新）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
NETR = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    lab = r[0].value
    if lab in ("  其中 CPO", "  其中可插拔", "AI IT 資本支出（加速器 × 每顆 IT capex）", "scale-out 光模組"):
        NETR[lab.strip()] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETR) == 4 and all(v is not None for k in NETR for v in NETR[k][0]), NETR.keys()
# 網通 CPO 三情境（markets/network v0.2.3；Andy 2026-10-25 13:13：技術替代屬行業情境，A＝替代快、C＝慢，公司讀同字母）
NETS = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    for s_ in ("A", "B", "C"):
        for k_, l_ in (("mpl", f"  其中可插拔（情境 {s_}）"), ("mcp", f"  其中 CPO（情境 {s_}）")):
            if r[0].value == l_:
                NETS[(k_, s_)] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert len(NETS) == 6 and all(v is not None for k in NETS for v in NETS[k][0]), NETS.keys()
NETR["其中可插拔"] = NETS[("mpl", "A")]; NETR["其中 CPO"] = NETS[("mcp", "A")]   # 情境 A 列（數值＝主表「其中可插拔」「其中 CPO」）
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。公司金額：十億人民幣（RMB B）；市場：十億美元（US$B）；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝深交所定期報告與公告（經財經媒體轉載）；公司不給營收指引；客戶名稱、800G／1.6T 份額只有次級來源"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
H1 = "https://finance.sina.com.cn/stock/hkstock/ggscyd/2026-08-21/doc-iniparan1339641.shtml"
H1M = "https://news.10jqka.com.cn/field/20260822/679192522.shtml"
FY25 = "https://finance.eastmoney.com/a/202603313690297170.html"
K36 = "https://eu.36kr.com/en/p/3917474184195719"
HCT = "https://hellochinatech.com/p/zhongji-innolight-optical-transceiver"
CONS = "https://basic.10jqka.com.cn/300308/worth.html"
LC = "https://lightcounting.com/newsletter/en/may-2026-optical-vendor-landscape-376"
AJ = "https://pdf.dfcfw.com/pdf/H3_AP202610121830372080_1.pdf"

sec("A. AI 光模組市場（讀 markets/network 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mpl_{y}", f"AI 可插拔光模組市場 20{y}", "US$B", None, NETR["其中可插拔"][0][j], None, "Derived", NETR["其中可插拔"][1],
        "加速器顆數 × 每顆光學內容（scale-out 為主；2029 含 scale-up 光學）− CPO" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"mcp_{y}", f"AI CPO 光學市場 20{y}", "US$B", None, NETR["其中 CPO"][0][j], None, "Derived", NETR["其中 CPO"][1], "共同封裝光學；網通模型的 CPO 滲透路徑" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][0][j], None, "Derived", NETR["AI IT 資本支出（加速器 × 每顆 IT capex）"][1], "Recon 用：中際旭創 AI 營收 ÷ AI IT 資本支出" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、Arista 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")

sec("B. 中際旭創份額與單價（本專案為權威：網通內容與廠商份額，規則第 13 條）")
inp("rel_27", "可插拔份額 2027 ÷ 2026（相對值）", "x", 0.85, 0.95, 1.05, "Assumed", "LightCounting 2025 排名：中際旭創第一（US$5.3B，+61%）、新易盛第二（US$3.5B，+189%，首次超過 Coherent）（Interested-party）", HCT + "：800G 約 35–40%、1.6T 約 50–70%（次級來源）")
inp("rel_28", "可插拔份額 2028 ÷ 2026", "x", 0.70, 0.90, 1.05, "Assumed", "同上；新易盛、Coherent 搶 1.6T 份額", "美國 FCC 26-50（2026-07）光模組目前只列為揭露元件；1260H 清單（2026-06）（鉅亨轉引花旗）")
inp("rel_29", "可插拔份額 2029 ÷ 2026", "x", 0.60, 0.85, 1.05, "Assumed", "同上", "—")
inp("us_27", "美國限制情境：份額相對值 2027（Sens 用）", "x", None, 0.90, None, "Assumed", "https://hao.cnyes.com/post/262773：花旗認為短期嚴格執行機率低；中國廠供美國大型雲端高速光模組約 60–70%", "泰國產能不必然豁免（美國價值含量門檻 65%，2029 起 75%）")
inp("us_28", "美國限制情境：份額相對值 2028（Sens 用）", "x", None, 0.55, None, "Assumed", "同上", "—")
inp("us_29", "美國限制情境：份額相對值 2029（Sens 用）", "x", None, 0.50, None, "Assumed", "同上", "—")
inp("cpo_27", *SIL.row9("cposh_innolight_27")[1:])
inp("cpo_28", *SIL.row9("cposh_innolight_28")[1:])
inp("cpo_29", *SIL.row9("cposh_innolight_29")[1:])
inp("pxr_27", "單價係數 2027（相對網通模型每顆內容；基準 1）", "x", 0.90, 1.0, 1.05, "Assumed", "網通模型每顆光學內容已含世代單價；本係數只在 Sens 動（降價快於模型）", "公司 10-23：價格體系健康、競爭有序（Interested-party）")
inp("pxr_28", "單價係數 2028", "x", 0.80, 1.0, 1.05, "Assumed", "同上", "—")
inp("pxr_29", "單價係數 2029", "x", 0.72, 1.0, 1.05, "Assumed", "同上", "—")
inp("fx", "人民幣兌美元（RMB／USD）", "RMB", 6.9, 7.1, 7.3, "Assumed", FY25 + "：人民幣兌美元每變動 4%，稅前利潤約變動 6.45 億元（公司）", "海外營收占 90.6%（FY2025，公司）")

sec("C. 其他營收（國內、電信、100G 以下、元件；網通模型未涵蓋）")
inp("oth_s", "其他營收 ÷ 總營收（FY2025、FY2026）", "%", 0.05, 0.08, 0.12, "Assumed", FY25 + "：國內營收 36.03 億元（9.4%），含國內 AI 的 400G／800G", "光模組以外產品占 1.0%（搜狐轉載公司資料）")
inp("g_oth", "其他營收年增（FY2027 起）", "%", 0.0, 0.10, 0.20, "Assumed", "國內營收 FY2025 +14.5%（公司）", "—")

sec("D. 毛利率")
inp("gm_h2", "2H26 毛利率", "%", 0.45, 0.47, 0.49, "Assumed", H1M + "：1H26 光模組毛利率 46.59%（+6.63pt）", HCT + "：1Q26 46%（歷史新高）→ 2Q26 約 47%（推導）")
inp("gm_27", "毛利率 2027", "%", 0.43, 0.47, 0.49, "Assumed", "1.6T、矽光占比上升（上推）vs 降價（下壓）", AJ + "：光學元件與雷射約占成本 25–35%、DSP 等電子 20–28%（Singularity Research，次級來源）")
inp("gm_28", "毛利率 2028", "%", 0.40, 0.46, 0.48, "Assumed", "同上；FY2024 光模組 34.65%、FY2025 42.61%（公司）", "—")
inp("gm_29", "毛利率 2029", "%", 0.36, 0.45, 0.47, "Assumed", "同上", "—")

sec("E. FY2026 錨點（1H26 實績＋2H26 假設）")
inp("h1_rev", "1H26 營收", "RMB B", None, 41.778, None, "Verified", H1 + "：417.78 億元（+182.49%）", H1M + "：同值；光模組 413.55 億元")
inp("h1_gm", "1H26 毛利率（光模組，占營收 99%）", "%", None, 0.4659, None, "Verified", "公司 1H26 報告（同花順轉載，網址見 Sources）：46.59%", "全公司毛利率未揭露：未找到（以光模組代用）")
inp("h1_ni", "1H26 歸母淨利", "RMB B", None, 13.651, None, "Verified", H1 + "：136.51 億元（+241.70%）", H1M + "：同值")
inp("h1_eps", "1H26 基本 EPS", "RMB", None, 12.31, None, "Verified", H1 + "：12.31 元", "—")
inp("q1_rev", "1Q26 營收", "RMB B", None, 19.5, None, "Interested-party", "Hello China Tech（網址見 Sources）：Rmb 19.5bn（+192%）", "一季報原文：未找到")
inp("q_g", "3Q26、4Q26 營收季增", "%", 0.05, 0.15, 0.25, "Assumed", "2Q26 ÷ 1Q26 − 1 ≈ +14%（推導）；公司 10-23：訂單覆蓋到 2027、部分客戶給出 2028 指示", "共識 FY2026 101.5 隱含 2H26 ≈ 59.7（季增約 +20%，推導）")
inp("tx", "有效稅率", "%", 0.12, 0.15, 0.18, "Assumed", "高新技術企業 15%；海外子公司（泰國等）不同：未找到", "—")
inp("attr", "歸母淨利 ÷ 淨利（含少數股權）", "x", 0.90, 0.932, 0.97, "Derived", "東方財富轉載 FY2025 年報（網址見 Sources）：歸母 107.97 ÷ 淨利 115.8 億元", "1H26 少數股權：未找到")
inp("sh_post", "H 股發行後總股數", "十億股", None, "=722.447/613.33", None, "Derived", "https://www.sohu.com/a/1091020082_122014422：總市值 7,224.47 億 ÷ 股價 613.33（2026-10-22 盤中）", K36 + "：H 股 2026-07 發行，每股 HK$980")
inp("q_sh_h1", "1H26 加權股數", "十億股", None, "=h1_ni/h1_eps", None, "Derived", "歸母淨利 ÷ 基本 EPS", "—")

sec("F. FY2025 實績")
inp("a25_rev_in", "FY2025 營收", "RMB B", None, 38.24, None, "Verified", K36 + "：382.4 億元（+60.25%）", FY25 + "：同值")
inp("a25_gm_in", "FY2025 毛利率（光模組代用）", "%", None, 0.4261, None, "Verified", "東方財富轉載 FY2025 年報（網址見 Sources）：光模組 42.61%（+7.96pt）", "全公司毛利率：未找到")
inp("a25_ni_in", "FY2025 歸母淨利", "RMB B", None, 10.797, None, "Verified", K36 + "：107.97 億元（+108.78%）", FY25 + "：同值")
inp("a25_sh_in", "FY2025 股數（近似）", "十億股", None, "=q_sh_h1", None, "Derived", "FY2025 年報 EPS：未找到；以 1H26 加權股數近似", "—")

sec("G. 營業費用、股數（FY2027 起）")
inp("phi", "營業費用（含減值與其他）成長對營收成長的彈性", "x", 0.40, "=Q!B26", 0.80, "Derived", "FY2026 錨點營業費用年增 ÷ 營收年增（Q 頁）", "FY2025 銷售＋管理＋研發＋財務費用 27.86 億元、銷售費用率 0.59%（公司）")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.01, 0.0, 0.02, "Assumed", "無", "Sens 用 +2 個百分點")
inp("sh_g", "股數年變動", "%", -0.01, 0.0, 0.01, "Assumed", "H 股發行後無已知增發；有回購（9 月 806.93–944.68 元成交）", "—")
inp("mf", "AI 市場係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年可插拔光模組 +84%", "±20%")

sec("H. 共識、市場資料（只作 Recon 對帳，不作校準）")
for y, rv, ni, eps in ((26, 101.466, 33.645, 29.28), (27, 186.360, 62.985, 54.74), (28, 273.291, 91.017, 79.16)):
    inp(f"cons_rev{y}", f"共識 FY20{y} 營收", "RMB B", None, rv, None, "Interested-party", CONS + "（15 家，2026-10-23）", "區間見同頁")
    inp(f"cons_ni{y}", f"共識 FY20{y} 歸母淨利", "RMB B", None, ni, None, "Interested-party", CONS + "（24 家）", "—")
    inp(f"cons_eps{y}", f"共識 FY20{y} EPS", "RMB", None, eps, None, "Interested-party", CONS + "（24 家）", "—")
inp("gs_ni27", "高盛 FY2027 歸母淨利（單一券商）", "RMB B", None, 86.9, None, "Interested-party", "https://www.21jingji.com/article/20261023/c6f736ac52a4cda8658a9afc828059c5.html：高盛（9 月）2026、2027 約 404、869 億元，目標價 2,645 元", "H 股聯席保薦人之一（利害關係）")
inp("lc_dc26", "LightCounting 2026 全球數通光模組市場", "US$B", None, 22.8, None, "Interested-party", AJ + "：約 228 億美元（引 LightCounting）", "網通模型 2026 可插拔 AI 光學較高（見 A 節）")
inp("sh800", "800G 份額（次級來源，中點）", "%", 0.35, 0.375, 0.40, "Interested-party", HCT + "：35–40%", "Photon Cap：中際旭創＋新易盛約占 NVIDIA 800G 60%")
inp("sh16", "1.6T 份額（次級來源，中點）", "%", 0.50, 0.60, 0.70, "Interested-party", HCT + "：50–70%", K36 + "：1.6T 全球份額 50% 以上（未經證實的公開資訊）")
inp("pe_fair_andy", "Andy 先驗：合理 P/E", "x", None, 30, None, "Assumed", "先驗_Andy_中際旭創_20261024 Q8b（Drive；不在本 repo）", "只作對照")

sec("I. 價格與估值")
inp("px", "股價（2026-10-23 收盤，A 股）", "RMB", None, 612.63, None, "Interested-party", "https://www.etnet.com.hk/www/sc/ashares/quote.php?code=300308：612.63（+3.82%；前收 590.11）", "https://www.21jingji.com/article/20261023/c6f736ac52a4cda8658a9afc828059c5.html：10-22 收盤 590.11")
inp("px_hi", "52 週高點", "RMB", None, 1416.88, None, "Interested-party", "https://www.google.com/finance/quote/300308:SHE?hl=zh：52 週區間 337.00–1,416.88", "—")
sec("J. 目標價情境權重（v0.1.1 新增；Val 頁第 0 節用）")
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")

sec("K. 目標價（DCF＋Comparable）與收斂法參數（v0.2 新增；黃底，Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；各市場同一值）", "DCF 與 Sens")
inp("tp_g", "長期成長率（DCF 永續；收斂法情境 B 的 FY2032 營收年增）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "—")
inp("tp_roe", "ROE（近三年平均；DCF 再投資率＝tp_g ÷ tp_roe）", "%", None, '=AVERAGE(0.1645,0.3064,0.4461)', None, 'Derived', 'stockanalysis 比率頁：FY2023 16.45%、FY2024 30.64%、FY2025 44.61%（GAAP 歸母口徑）', '單一來源；近 12 月 64.62%（同站）')
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（工作單 2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "收斂法：情境 A 於 FY2032 的營收年增（＝tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法；FY2030 起由 FY2029 年增線性收斂", "FY2030 起為收斂法，非驅動因子推導")
inp("fade_gB", "收斂法：情境 B 於 FY2032 的營收年增（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "收斂法：情境 C 於 FY2032 的營收年增（＝tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "收斂法：情境 A 的中段淨利率（＝情境 A FY2029 淨利率）", "%", None, None, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "收斂法：情境 B 的中段淨利率（＝FY2029 與十年平均的中點；十年平均以 FY2029 × 0.8 代替）", "%", None, None, None, "Assumed", "十年平均淨利率（2016–2025，與本模型同口徑）：未找到，依工作單以 FY2029 × 0.8 代替；本列＝（FY2029＋FY2029 × 0.8）÷ 2", "Andy 可改")
inp("fade_mC", "收斂法：情境 C 的中段淨利率（＝十年平均；以情境 C FY2029 × 0.8 代替）", "%", None, None, None, "Assumed", "十年平均淨利率：未找到，依工作單以情境 C FY2029 淨利率 × 0.8 代替", "Andy 可改")
inp("cons_tp", "市場共識目標價（賣方平均）", 'RMB', None, 1321, None, "Interested-party", 'stockanalysis（S&P Global 彙整）預測頁：14 位分析師平均 1,321、中位數 1,400、區間 430–1,800（2026-10-25 擷取）', '同站統計頁 1,392.11（17 位）；兩頁家數不同，取預測頁；多數目標價訂於 10 月股價下跌前')
inp("cons_feps", '共識 12 個月後前瞻 EPS（2/12 × FY2027＋10/12 × FY2028）', 'RMB', None, '=2/12*cons_eps27+10/12*cons_eps28', None, "Derived", 'In 頁同花順共識 FY2027、FY2028 EPS（與本模型前瞻 EPS 同期間）', "共識只放 Recon 與本列，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億人民幣；EPS 為人民幣；股數為十億股）。藍字＝輸入；黑字＝公式。口徑：歸母淨利；毛利率以光模組（占營收 99%）代用"
Q["A2"] = "FY2026 錨點＝1H26 實績＋2H26（3Q、4Q 營收依季增假設；毛利率 gm_h2；營業費用率同 1H26）。營業費用、減值與其他為殘差：毛利 − 歸母淨利 ÷［（1 − 稅率）× 歸母比例］。Q3 季報（10-31 前）公布後以實績取代 3Q"
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（假設）", "4Q26（假設）", "2H26 合計", "FY2026 錨點", "標記", "說明"])
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", note=""):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, note)
qrow(4, "營收", {"B": "=a25_rev_in", "C": "=h1_rev", "D": "=(h1_rev-q1_rev)*(1+q_g)", "E": "=D4*(1+q_g)", "F": "=D4+E4", "G": "=C4+F4"}, tag="Verified／Assumed（3Q、4Q）", note="2Q26＝1H26 − 1Q26；3Q＝2Q ×（1＋季增）")
qrow(5, "毛利", {"B": "=B4*a25_gm_in", "C": "=C4*h1_gm", "D": "=D4*gm_h2", "E": "=E4*gm_h2", "F": "=D5+E5", "G": "=C5+F5"}, tag="Derived")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEFG"}, fmt="0.0%", tag="Derived")
qrow(7, "稅後歸母係數 ＝（1 − 稅率）× 歸母比例", {c: "=(1-tx)*attr" for c in "BCDE"}, fmt="0.000", tag="Assumed／Derived")
qrow(8, "營業費用、減值與其他（淨）", {"B": "=B5-B10/B7", "C": "=C5-C10/C7", "D": "=D4*C9", "E": "=E4*C9", "F": "=D8+E8", "G": "=C8+F8"}, tag="Derived（殘差）", note="FY2025 殘差 ÷ 營收約 7%，與公司四項費用合計 27.86 億元（7.3%）相近")
qrow(9, "營業費用率", {c: f"={c}8/{c}4" for c in "BCDEFG"}, fmt="0.0%", tag="Derived")
qrow(10, "歸母淨利", {"B": "=a25_ni_in", "C": "=h1_ni", "D": "=(D5-D8)*D7", "E": "=(E5-E8)*E7", "F": "=D10+E10", "G": "=C10+F10"}, tag="Verified／Derived（3Q、4Q）")
qrow(11, "淨利率", {c: f"={c}10/{c}4" for c in "BCDEFG"}, fmt="0.0%", tag="Derived")
qrow(12, "股數（十億股）", {"B": "=a25_sh_in", "C": "=q_sh_h1", "D": "=sh_post", "E": "=sh_post", "G": "=(C12*2+D12+E12)/4"}, tag="Derived", note="FY2026＝1H 加權 × ½ ＋ H 股發行後 × ½（近似）")
qrow(13, "EPS", {"B": "=B10/B12", "C": "=C10/C12", "D": "=D10/D12", "E": "=E10/E12", "G": "=G10/G12"}, fmt="0.00", tag="Derived", note="1H26 公司揭露 12.31")
qrow(14, "其他營收（國內、電信等；＝營收 × oth_s）", {"B": "=B4*oth_s", "G": "=G4*oth_s"}, tag="Assumed")
qrow(15, "AI 營收（＝營收 − 其他）", {"B": "=B4-B14", "G": "=G4-G14"}, tag="Derived")
qrow(16, "AI 營收（US$B ＝ ÷ fx）", {"B": "=B15/fx", "G": "=G15/fx"}, tag="Derived")
Q.cell(21, 1, "FY2026 錨點與共識的距離（參考；校準未使用共識）").font = BOLD
Q.cell(22, 1, "FY2026 營收 ÷ 共識 − 1"); Q["B22"] = "=G4/cons_rev26-1"
Q.cell(23, 1, "FY2026 歸母淨利 ÷ 共識 − 1"); Q["B23"] = "=G10/cons_ni26-1"
Q.cell(24, 1, "FY2026 營收年增"); Q["B24"] = "=G4/B4-1"
Q.cell(25, 1, "FY2026 營業費用年增"); Q["B25"] = "=G8/B8-1"
Q.cell(26, 1, "營業費用彈性（營業費用年增 ÷ 營收年增；In 頁 phi 的依據）"); Q["B26"] = "=B25/B24"
Q.cell(27, 1, "1H26 EPS（模型）vs 公司揭露 12.31"); Q["B27"] = "=C13/h1_eps-1"
for rr in range(22, 28): Q[f"B{rr}"].number_format = "0.0%" if rr != 26 else "0.00"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 8), ("ni", 10), ("sh", 12), ("eps", 13), ("oth", 14), ("ai", 15), ("aiu", 16)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 52
for c in "BCDEFG": Q.column_dimensions[c].width = 15
Q.column_dimensions["H"].width = 26; Q.column_dimensions["I"].width = 80
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
mk = lambda b, y, s: f"{b}_{y % 100}" if s == "A" else f"{b}{s}_{y % 100}"   # 網通 CPO 三情境：情境 A 讀 A 節，B、C 讀同字母列
DEFP = dict(scn="A", rel="rel", cpo="cpo", pxr="pxr", gm="gm", mf="1", fx="fx", gn="g_oth", opxa="0")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    F = BU[1:]
    sec_("A. AI 營收（D 欄＝FY2026 校準年；E–G＝FY2027–FY2029）")
    row("mpl", "AI 可插拔光模組市場 × 情境顆數係數", "US$B", {c_(y): f"={mk('mpl', y, s)}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("mcp", "AI CPO 光學市場 × 情境顆數係數", "US$B", {c_(y): f"={mk('mcp', y, s)}*{umx(y, s)}" for y in BU})
    row("fxr", "人民幣兌美元", "RMB", {c_(y): ("=fx" if y == 2026 else f"={P['fx']}") for y in BU}, fmt="0.00",
        note="FY2026 校準用基準匯率；Sens 只動 FY2027 起")
    row("sh0", "中際旭創可插拔有效份額（FY2026 校準）＝ AI 營收（美元）÷ 可插拔市場", "%", {"D": f"=a26_aiu/D{R['mpl']}"}, fmt="0.0%",
        note="Derived：FY2026 CPO 營收視為 0；FY2027 起乘相對值")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in F}}, fmt="0.00")
    row("sh", "可插拔有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("pxr", "單價係數（相對網通模型）", "x", {"D": "=1", **{c_(y): f"={P['pxr']}_{y % 100}" for y in F}}, fmt="0.00")
    row("cps", "中際旭創占 CPO 光學金額", "%", {"D": "=0", **{c_(y): f"={P['cpo']}_{y % 100}" for y in F}}, fmt="0.0%")
    row("aiu", "AI 營收（美元）", "US$B", {"D": "=a26_aiu", **{c_(y): f"={c_(y)}{R['mpl']}*{c_(y)}{R['sh']}*{c_(y)}{R['pxr']}+{c_(y)}{R['mcp']}*{c_(y)}{R['cps']}" for y in F}})
    row("ai", "AI 營收", "RMB B", {"C": "=a25_ai", **{c_(y): f"={c_(y)}{R['aiu']}*{c_(y)}{R['fxr']}" for y in BU}})
    sec_("B. 其他營收與合計")
    row("oth", "其他營收（國內、電信等）", "RMB B", {"C": "=a25_oth", "D": "=a26_oth", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']})" for y in F}})
    row("rev", "營收", "RMB B", {"C": "=a25_rev", **{c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['oth']}" for y in BU}})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    sec_("C. 損益")
    row("gm", "毛利率", "%", {"C": "=a25_gp/a25_rev", "D": "=a26_gp/a26_rev", **{c_(y): f"={P['gm']}_{y % 100}" for y in F}}, fmt="0.0%",
        note="FY2027 起為輸入路徑（1.6T 與矽光上推 vs 降價下壓）")
    row("gp", "毛利", "RMB B", {c_(y): f"={c_(y)}{R['rev']}*{c_(y)}{R['gm']}" for y in FYS})
    row("opb", "營業費用、減值與其他（基準路徑）", "RMB B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,phi*{c_(y)}{R['g']}))" for y in F}})
    row("opx", "營業費用、減值與其他", "RMB B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in F}})
    row("pre", "稅前利益（＝毛利 − 營業費用等）", "RMB B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS})
    row("prem", "稅前利益率", "%", {c_(y): f"={c_(y)}{R['pre']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("ni", "歸母淨利", "RMB B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"={c_(y)}{R['pre']}*(1-tx)*attr" for y in F}})
    row("nim", "淨利率", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", "E": "=sh_post*(1+sh_g)", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in (2028, 2029)}}, fmt="0.000")
    row("eps", "EPS", "RMB", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 中際旭創損益（RMB B）。FY2025A 實績、FY2026 錨點（1H26 實績＋2H26 假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026 三情境相同；FY2027 起三情境讀不同的網通 CPO 情境（A 替代快、C 慢），顆數係數從 CY2028 開始"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）；網通 CPO 情境 A（替代快）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85；網通 CPO 情境 B", "C": "情境 C：CY2028–29 顆數 ×0.6；網通 CPO 情境 C（替代慢）"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "RMB B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 營收", "RMB B", "#,##0.0"), ("oth", "  其他營收", "RMB B", "#,##0.0"),
                                ("aish", "  AI ÷ 營收", "%", "0.0%"), ("gm", "毛利率", "%", "0.0%"), ("nim", "淨利率", "%", "0.0%"), ("ni", "歸母淨利", "RMB B", "#,##0.0"),
                                ("eps", "EPS", "RMB", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型營收、AI 營收、歸母淨利 − 錨點（校準恆等式）").font = BOLD; sr += 1
PL.cell(sr, 1, "  情境 A：營收差（RMB B）"); PL[f"D{sr}"] = f"=D{PB['A']['rev']}-a26_rev"; PL[f"D{sr}"].number_format = "0.0000"; CHK_ROW = sr; sr += 1
PL.cell(sr, 1, "  情境 A：AI 營收差（RMB B）"); PL[f"D{sr}"] = f"=D{PB['A']['ai']}-a26_ai"; PL[f"D{sr}"].number_format = "0.0000"; sr += 1
PL.cell(sr, 1, "  情境 A：稅前利益 ×（1 − 稅率）× 歸母比例 − 歸母淨利錨點"); PL[f"D{sr}"] = f"=D{PB['A']['pre']}*(1-tx)*attr-a26_ni"; PL[f"D{sr}"].number_format = "0.0000"
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：市場與份額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 中際旭創在 AI 光模組的份額（本專案權威，規則第 13 條）。市場取自 markets/network（加速器顆數 × 每顆光學內容；美元）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("mpl", "AI 可插拔光模組市場（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mpl']}", "#,##0.0", "markets/network Mkt 頁"),
        ("mcp", "AI CPO 光學市場（情境 A）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mcp']}", "#,##0.0", "同上"),
        ("cpos", "CPO ÷（可插拔＋CPO）", "%", lambda y: f"=PL!{FC[y]}{A_['mcp']}/(PL!{FC[y]}{A_['mpl']}+PL!{FC[y]}{A_['mcp']})", "0.0%", "網通模型的 CPO 滲透率"),
        ("sh", "中際旭創可插拔有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以 1H26 實績＋2H26 假設校準"),
        ("aiu", "中際旭創 AI 營收（情境 A；美元）", "US$B", lambda y: f"=PL!{FC[y]}{A_['aiu']}", "#,##0.00", ""),
        ("oth", "其他廠商（新易盛、Coherent 等）＝ 可插拔＋CPO − 中際旭創", "US$B", lambda y: f"=PL!{FC[y]}{A_['mpl']}+PL!{FC[y]}{A_['mcp']}-PL!{FC[y]}{A_['aiu']}", "#,##0.0", ""),
        ("itr", "中際旭創 AI 營收 ÷ AI IT 資本支出", "%", lambda y: f"=PL!{FC[y]}{A_['aiu']}/itc_{y % 100}", "0.00%", "markets/network 的 AI IT 資本支出（加速器 × 每顆 IT capex）")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：FY2026 有效份額落在次級來源 800G 約 35–40%、1.6T 約 50–70% 之間；份額相對值與 CPO 份額是 Assumed，是最需要每季檢查的輸入。").font = BOLD
UN.column_dimensions["A"].width = 64; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（人民幣；A 股）。目標價＝DCF 與 Comparable 加權（第 0 節）；結論另列「現價 ÷ FY2027、FY2028 EPS（模型）」、市場隱含 EPS 與十年 P/E 規則值（參考）。十年 P/E 只用 2018–2025（2017 年併購旭創完成前為中際裝備，不可比）"
V["A2"] = "循環位置：中際旭創不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照 Arista；FY(t)＝CY(t)，不平移）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-23 收盤）"); V["B6"] = "=px"; V["B6"].number_format = "0.00"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "現價 ÷ 共識未來 12 月 EPS（2/12 × FY2026＋10/12 × FY2027）"); V["B11"] = "=px/(cons_eps26*2/12+cons_eps27*10/12)"; V["B11"].number_format = "0.0"
V.cell(12, 1, "現價 ÷ 52 週高點 − 1"); V["B12"] = "=px/px_hi-1"; V["B12"].number_format = "0%"
V.cell(13, 1, "現價 ÷ 近 12 月 EPS（2H25＋1H26；推導）"); V["B13"] = "=px/((a25_ni-h1_ni/3.417+h1_ni)/sh_post)"; V["B13"].number_format = "0.0"
V["E8"] = "FY2027 三情境只差在網通 CPO 情境（顆數係數從 CY2028 開始）"; V["E13"] = "1H25 歸母淨利＝1H26 ÷（1＋241.7%）"
V.cell(15, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, 16, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 17):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
V.cell(21, 1, "3. 規則值隱含股價（人民幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, 22, ["情境", "FY2027 EPS", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = 23 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=D{17 + j}").number_format = "0.0"
    V.cell(i, 4, f"=B{i}*C{i}").number_format = "#,##0"
    V.cell(i, 5, f"=D{i}/px-1").number_format = "0%"
    V.cell(i, 6, f"=px/B{i}").number_format = "0.0"
    for c in range(1, 7): V.cell(i, c).fill = KEY
    VAL[s] = i
V.cell(26, 1, "4. 核心問題：市場隱含的 FY2028 EPS（＝現價 ÷ 假設倍數；約 2027 年底的價格看下一年 EPS）與模型的距離").font = BOLD
hdr(V, 27, ["假設倍數（÷ FY2028 EPS）", "市場隱含 FY2028 EPS", "模型 FY2028 EPS（A）", "隱含 ÷ 模型 A − 1", "模型 FY2028 EPS（C）", "隱含 ÷ 模型 C − 1"])
vr = 28; IMPL = {}
for m in (10, 15, 20, "=pe_peak", "=pe_fair_andy"):
    V.cell(vr, 1, m).number_format = "0.0"
    if not isinstance(m, str): V.cell(vr, 1).font = BLUE
    V.cell(vr, 2, f"=px/A{vr}").number_format = "0.00"
    V.cell(vr, 3, f"=PL!F{SUMR[('A', 'eps')]}").number_format = "0.00"
    V.cell(vr, 4, f"=B{vr}/C{vr}-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[('C', 'eps')]}").number_format = "0.00"
    V.cell(vr, 6, f"=B{vr}/E{vr}-1").number_format = "0%"
    IMPL[m] = vr; vr += 1
V.cell(vr, 1, "讀法：第四列＝十年 P/E 25 分位（見頂倍數）、第五列＝Andy 合理 P/E。若市場對 FY2028 用見頂倍數，現價隱含的 FY2028 EPS 遠低於模型——要成立，需要份額、單價或 CPO 替代其中至少一項大幅不利（見 Sens）。")
vr += 2
V.cell(vr, 1, "5. 對照：以 Andy 合理 P/E 30x 計算的隱含價（只作對照，非規則值）").font = BOLD; vr += 1
hdr(V, vr, ["情境", "FY2027 EPS", "× 30x", "相對現價", "FY2028 EPS", "× 30x", "相對現價"]); vr += 1
IMP = {}
for s in SC:
    V.cell(vr, 1, s)
    V.cell(vr, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 3, f"=B{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 4, f"=C{vr}/px-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[(s, 'eps')]}").number_format = "0.00"; V.cell(vr, 6, f"=E{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
    IMP[s] = vr; vr += 1
assert vr < 50, vr
PE_CMC = [None, None, 29.0, 98.4, 56.4, 49.5, 24.4, 56.0, 27.7, 64.5]   # companiesmarketcap 年底 P/E；2016–2017 為中際裝備（併購前），不用
block_end = VR.add_rule_block(V, 50, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2018–2025 共 8 年；2016–2017 為併購前的中際裝備，不用）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=B{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=B{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=B{VAL['A']}"}]}],
    "#,##0.00", "規則（Andy 2026-10-13）：合理倍數＝年底 P/E 依循環位置取分位。只有一個來源、8 年（第二來源：未找到）。歷史倍數為落後 P/E，且 2019、2025 年為高成長期（98x、65x），規則值可能高估成長放緩後的倍數；前瞻 P/E 歷史待 Bloomberg（抓取單已含 300308）。")
rule_rows = [rw for rw in range(50, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("份額流失（低）", "份額相對值 2027 0.85、2028 0.70、2029 0.60（新易盛、Coherent 搶 1.6T）", dict(rel="RLO")),
         ("份額持平（高）", "份額相對值 1.05", dict(rel="RHI")),
         ("美國限制", "份額相對值 2027 0.90、2028 0.55、2029 0.50（Covered List 擴及光模組）", dict(rel="us")),
         ("降價快於模型", "單價係數 2027 0.90、2028 0.80、2029 0.72", dict(pxr="PXLO")),
         ("毛利率 低（價格戰）", "毛利率 2027 43%、2028 40%、2029 36%", dict(gm="GMLO")), ("毛利率 高", "毛利率 2027 49%、2028 48%、2029 47%", dict(gm="GMHI")),
         ("CPO 份額 0", "中際旭創不參與 CPO 光學", dict(cpo="CPLO")), ("CPO 份額 高", "2027 30%、2028 35%、2029 40%", dict(cpo="CPHI")),
         ("AI 市場 ×0.8", "FY2027 起可插拔光模組市場 ×0.8", dict(mf="mf_lo")), ("AI 市場 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("人民幣升值（6.9）", "FY2027 起 RMB／USD 6.9", dict(fx="fx_lo")), ("人民幣貶值（7.3）", "FY2027 起 7.3", dict(fx="fx_hi")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi"))]
for nm, base, suf in (("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("PXLO", "pxr", "_lo"), ("GMLO", "gm", "_lo"), ("GMHI", "gm", "_hi"), ("CPLO", "cpo", "_lo"), ("CPHI", "cpo", "_hi")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}_{y}"] = DefinedName(f"{nm}_{y}", attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "ai", 2027, "#,##0.0"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "ai", 2028, "#,##0.0"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85；CPO 情境 B", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6；CPO 情境 C", PB["C"], "PL!")
assert nn + 3 < 40
SENS_N = len(cases)
SE.cell(nn + 3, 1, "讀法：FY2027 三情境只差在網通 CPO 情境（A 替代快、C 慢）；顆數差異從 FY2028 起。「美國限制」「份額流失」「降價」「毛利率」四案例對應市場隱含衰退的三條路徑（份額、單價、CPO 替代）；CPO 份額案例顯示 CPO 金額在 2029 年前仍小。")
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識、券商、研調份額。只作事後對照，不為對上數字而調參數。共識區間很寬（FY2027 EPS 35.4–75.7），差距需併看區間"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與券商（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "同花順（15 家）", "錨點＝1H26 實績＋2H26 季增假設")
rec("FY2026 歸母淨利 vs 共識", "=cons_ni26", pl("ni", 2026), "同花順（24 家）", "")
rec("FY2027 營收 vs 共識", "=cons_rev27", pl("rev", 2027), "同花順（15 家；113.4–266.4）", "")
rec("FY2027 歸母淨利 vs 共識", "=cons_ni27", pl("ni", 2027), "同花順（24 家；39.4–89.2）", "")
rec("FY2027 歸母淨利 vs 高盛", "=gs_ni27", pl("ni", 2027), "21 經濟網轉引高盛（9 月）", "高盛為 H 股聯席保薦人")
rec("FY2027 EPS vs 共識", "=cons_eps27", pl("eps", 2027), "同花順（24 家；35.37–75.71）", "")
rec("FY2028 營收 vs 共識", "=cons_rev28", pl("rev", 2028), "同花順", "")
rec("FY2028 EPS vs 共識", "=cons_eps28", pl("eps", 2028), "同花順（55.57–116.11）", "")
RC.cell(rr_, 1, "B. 份額與市場（先驗比對的機制檢查）").font = BOLD; rr_ += 1
rec("FY2026 可插拔有效份額 vs 800G 份額（次級來源中點）", "=sh800", lambda s: f"=PL!D{PB[s]['sh0']}", "Hello China Tech", "有效份額含 1.6T（份額較高）與國內以外的全部 AI 光學", fmt="0.0%", pt=True)
rec("FY2026 可插拔有效份額 vs 1.6T 份額（次級來源中點）", "=sh16", lambda s: f"=PL!D{PB[s]['sh0']}", "Hello China Tech", "", fmt="0.0%", pt=True)
rec("FY2026 AI 營收（美元）vs LightCounting 2026 數通光模組市場", "=lc_dc26", pl("aiu", 2026), "LightCounting（愛建證券轉引）", "比值＝隱含份額；LightCounting 預測可能偏保守", fmt="#,##0.0")
rec("FY2027 中際旭創 AI 營收 ÷ AI IT 資本支出", None, lambda s: f"=PL!E{PB[s]['aiu']}/itc_27", "markets/network", "每 1 美元 AI IT 資本支出分到的光模組營收", fmt="0.00%")
rec("FY2028 可插拔有效份額", None, pl("sh", 2028), "—", "", fmt="0.0%")
RC.cell(rr_, 1, "C. 校準").font = BOLD; rr_ += 1
rec("1H26 EPS（模型）vs 公司揭露 12.31", "=h1_eps", lambda s: "=Q!C13", "公司", "股數以淨利 ÷ EPS 推導，恆等", fmt="0.00")
rec("FY2026 營業費用率（Derived）", None, lambda s: "=Q!G9", "Q 頁殘差", "", fmt="0.0%")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以中際旭創 AI 營收 ÷ 部署 GW 計算每 GW 營收"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.cell(rr_, 1, "3Q26 季報（10-31 前）：公布後以實績取代 Q 頁 3Q 假設"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝深交所定期報告與公告，經財經媒體轉載；原文 PDF 待人工複核）"; SO["A1"].font = BOLD
srcs = [("公司（轉載）：1H26 實績（新浪財經）", H1),
        ("公司（轉載）：1H26 光模組毛利率 46.59%、海外營收（同花順）", H1M),
        ("公司（轉載）：FY2025 光模組毛利率 42.61%、研發與財務費用、匯率敏感度（東方財富）", FY25),
        ("二手：FY2025 實績、客戶集中度（前五大 76%、最大 24%）、1.6T 占比、H 股發行（36氪／鳳凰週刊）", K36),
        ("二手：客戶名稱、800G／1.6T 份額、1Q26 營收與毛利率、1260H（Hello China Tech；付費牆，只讀到預覽）", HCT),
        ("公司（轉載）：10-23 經營公告（訂單到 2027、NPO／XPO 送樣）", "https://news.china.com/socialgd/10000169/20261023/49813907.html"),
        ("共識：同花順 F10 盈利預測（2026-10-23）", CONS),
        ("研調：LightCounting 2025 廠商排名", LC),
        ("券商：愛建證券 10-12 通信行業報告（成本結構、LightCounting 市場規模）", AJ),
        ("券商：花旗對 FCC 26-50 的解讀（鉅亨）", "https://hao.cnyes.com/post/262773"),
        ("價格：etnet（10-23 收盤）、21 經濟網（10-22 收盤、高盛預測）、Google 財經（52 週區間）", "https://www.etnet.com.hk/www/sc/ashares/quote.php?code=300308"),
        ("資料商：companiesmarketcap 年底 P/E", "https://companiesmarketcap.com/zhongji-innolight/pe-ratio/"),
        ("市場層：markets/network、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：3Q26 季報；全公司毛利率；客戶名稱原始出處；800G／1.6T 份額第一手來源；1H26 少數股權；前瞻 P/E 歷史（Bloomberg）；P/E 第二來源", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("中際旭創（300308.SZ／3308.HK）收支與估值模型 v0.2.2（AI 半導體 Project：網通次產業內複製——光模組）　2026-10-26", True),
 ("版本：v0.2.1（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）；v0.2（2026-10-25：三情境延伸到 FY2032、目標價＝DCF 與 Comparable 加權；v0.1 初版、v0.1.1 加規則值情境加權）。中際旭創在 AI 光模組的份額在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）。", False),
 ("模型命題：中際旭創的 AI 營收＝AI 可插拔光模組市場（本專案網通模型）× 中際旭創份額 × 單價係數＋CPO 光學 × 中際旭創 CPO 份額；其他營收（國內、電信）另依成長率推。股價自高點腰斬、只約 12 倍共識未來 12 月 EPS：市場定價像是 2027 年就是獲利高點。", False),
 ("核心問題：市場隱含的 FY2028 EPS 遠低於模型——要成立，份額（新易盛、美國限制）、單價（每 Gb 降價）、CPO 替代三條路徑各要不利到什麼程度？", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029 由驅動因子推導；FY2030–FY2032 為收斂法。口徑：中國會計準則；淨利＝歸母淨利；EPS 以 H 股發行後股數。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 營收：markets/network 的 AI 可插拔光模組市場 × 情境顆數係數 × 中際旭創有效份額（FY2026 以 1H26 實績＋2H26 假設校準）× 份額相對值（Assumed）× 單價係數（基準 1）＋ CPO 光學市場 × CPO 份額（Assumed）；乘匯率換成人民幣。", False),
 ("其他營收：FY2026 營收 × 8%（Assumed）× 成長率。毛利率：FY2027 起為輸入路徑（Assumed）。營業費用、減值與其他：彈性 phi（Derived）。稅率 15%、歸母比例 0.932（Derived）。", False),
 ("校準：FY2026 ＝ 1H26 實績（營收 41.78、毛利率 46.59%、歸母淨利 13.65）＋ 2H26（季增 15%、毛利率 47%，Assumed）。公司不給指引；不用共識（只在 Recon 對帳）。", False),
 ("估值：目標價（Val 第 0 節）＝0.5 × DCF（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，情境權重 0.3／0.5／0.2）＋0.5 × Comparable（A 股光通訊同業前瞻 P/E 中位數 × 12 個月後前瞻 EPS）。另列：現價 ÷ FY2027、FY2028 EPS；市場隱含 FY2028 EPS（Val 第 4 節）；十年 P/E 規則值（2018–2025，一個資料商來源）× FY2027 EPS；循環位置以本模型營收最高年判定。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司原文經財經媒體轉載；客戶名稱與 800G／1.6T 份額只有次級來源。", False),
 ("", False),
 ("已知限制：(1) 份額相對值、CPO 份額、毛利率路徑皆為 Assumed，是 FY2028 EPS 最敏感的輸入（Sens）；(2) FY2026 錨點的 2H26 為假設，3Q26 季報 10-31 前公布；(3) 網通模型 2027 年可插拔光模組 +84%，未設雷射、DSP 供給上限（Sens 用市場 ×0.8）；(4) 十年 P/E 只有一個來源、8 年，含高成長期溢價；(5) 毛利率以光模組代用（占營收 99%）；(6) R4 產業部署路徑未建（Recon 預留列）。", False)]
lines.append(("v0.2（2026-10-25）：Val 頁最前面新增「0. 目標價」（0.1 DCF、0.2 Comparable、0.3 結論），既有內容整體下移；原「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留。PL 三情境延伸到 FY2032（FY2030 起為收斂法，非驅動因子推導；灰底），說明欄移到 K 欄。In 頁新增 K 節（tp_r、tp_g、tp_roe、tp_w_dcf、收斂法參數、共識目標價）。Sens 加目標價參數列。FY2026–FY2029 既有數字不變。", False))
lines.append(("v0.1.1（2026-10-25）：Val 頁新增「0. 目標價與上檔空間」（放在 Val 既有內容之後，不位移既有儲存格）與 In 頁目標價情境權重 tp_wA／tp_wB／tp_wC（0.3／0.5／0.2，Assumed）。各情境隱含價＝第 3 節規則值；其餘儲存格不變。", False))
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ---------------- 舊法：規則值情境加權（v0.1.1 新增時稱「0. 目標價與上檔空間」；v0.2 改名保留，只作參考） ----------------
# 為不位移既有儲存格（與舊版 dist 逐格相同），本節放在 Val 既有內容之後；HTML 估值章第一個區塊即為本節
def add_target_price(V, imp, note, fmt, method, px_note):
    t0 = V.max_row + 2
    V.cell(t0, 1, "舊法：規則值情境加權（參考）").font = BOLD
    hdr(V, t0 + 1, ["項目", "數值", "", "", "說明"])
    V.cell(t0 + 2, 1, "規則值情境加權價（參考，非目標價）"); V[f"B{t0 + 2}"] = f"=tp_wA*B{t0 + 4}+tp_wB*B{t0 + 5}+tp_wC*B{t0 + 6}"; V[f"B{t0 + 2}"].number_format = fmt
    V[f"E{t0 + 2}"] = "權重 × 各情境隱含價"
    V.cell(t0 + 3, 1, "規則值情境加權價 ÷ 現價 − 1（參考）"); V[f"B{t0 + 3}"] = f"=B{t0 + 2}/px-1"; V[f"B{t0 + 3}"].number_format = "0.0%"; V[f"E{t0 + 3}"] = px_note
    for i, (s, lab) in enumerate((("A", "情境 A 隱含價（樂觀）"), ("B", "情境 B 隱含價（基準）"), ("C", "情境 C 隱含價（下檔）"))):
        rr = t0 + 4 + i
        V.cell(rr, 1, lab); V[f"B{rr}"] = imp[s]; V[f"B{rr}"].number_format = fmt; V[f"E{rr}"] = note[s]
    rr = t0 + 7
    V.cell(rr, 1, "情境權重 A／B／C"); V[f"B{rr}"] = "=tp_wA"; V[f"C{rr}"] = "=tp_wB"; V[f"D{rr}"] = "=tp_wC"
    for c in "BCD": V[f"{c}{rr}"].number_format = "0%"
    V[f"E{rr}"] = f'="合計 "&TEXT(B{rr}+C{rr}+D{rr},"0%")&"；In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"'
    V.cell(t0 + 8, 1, "舊法說明"); V[f"B{t0 + 8}"] = method
    return t0
TP0 = add_target_price(V, {s: f"=D{VAL[s]}" for s in SC},
    {s: f'="第 3 節規則值：FY2027 EPS × "&TEXT(C{VAL[s]},"0.0")&"x（"&C{17 + SC.index(s)}&"）"' for s in SC}, "#,##0",
    "P/E：各情境 FY2027 EPS × 十年（2016–2025 年底）P/E 依循環位置取分位（以本模型營收最高年判定：A、B 中段＝中位數，C 見頂＝25 分位）；倍數只有 companiesmarketcap 一個來源、8 年（2018–2025）。"
    "十年倍數多為 AI 前、公司規模較小時期的高倍數（2019 年約 98x），偏誤方向為高估；FY2027 三情境 EPS 只差在網通 CPO 情境，A、B 隱含價相近、C 另差在倍數；隱含價約為 FY2027 年底的價格、未折現，當作 12 個月目標價時偏高。",
    "現價＝In 頁 px（2026-10-23 收盤，A 股）")


# ================= v0.2（2026-10-25）：FY2030–FY2032 收斂法；目標價＝DCF 與 Comparable 加權 =================
# 工作單：Andy 2026-10-25 11:18／11:20（三情境延伸到 FY2032、收斂法）、11:28（目標價＝DCF＋Comparable）、11:53（tp_w_dcf 0.5；A 股與其他市場同業分開）
# 本節在其他工作表都建好之後執行：(1) PL 說明欄由 H 移到 K，H–J 放 FY2030–32 收斂年度；(2) Val 既有內容整體下移，最前面放「0. 目標價」；(3) Sens 加目標價參數列
import re as _re
from copy import copy as _copy
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
FCX = dict(FC); FCX.update({2030: "H", 2031: "I", 2032: "J"})
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
INR = {IN.cell(i, 1).value: i for i in range(4, IN.max_row + 1) if IN.cell(i, 1).value}
for _k, _i in INR.items():
    if str(_k).startswith(("tp_", "fade_")) or _k in ("cons_tp", "cons_feps"):
        for _j in range(1, 10): IN.cell(_i, _j).fill = KEY

def _move_notes(ws):
    """說明欄（第 8 欄，H）移到第 11 欄（K），讓出 H–J；表頭列補上收斂年度。"""
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 8)
        if c.value is None: continue
        assert not (isinstance(c.value, str) and c.value.startswith("=")), (ws.title, r, c.value)
        k = ws.cell(r, 11); k.value = c.value; k.font = _copy(c.font); k.fill = _copy(c.fill)
        c.value = None; c.font = Font(); c.fill = PatternFill()
        if str(ws.cell(r, 1).value or "") in ("項目", "情境／項目"):
            for j, y in zip((8, 9, 10), FADE):
                h = ws.cell(r, j, f"FY{y}（收斂）"); h.font = BOLD; h.fill = HDR
    ws.column_dimensions["K"].width = ws.column_dimensions["H"].width or 70
    for col in "HIJ": ws.column_dimensions[col].width = 14

def _fade(ws, R, s, eps_f):
    """FY2030–32：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到）；非驅動因子推導。"""
    for y in FADE:
        c, pc, k = FCX[y], FCX[y - 1], y - 2029
        fs = {"g": f"=$G{R['g']}+(fade_g{s}-$G{R['g']})*{k}/3", "rev": f"={pc}{R['rev']}*(1+{c}{R['g']})",
              "nim": f"=$G{R['nim']}+(fade_m{s}-$G{R['nim']})*{k}/3", "ni": f"={c}{R['rev']}*{c}{R['nim']}",
              "sh_": f"={pc}{R['sh_']}*(1+sh_g)", "eps": "=" + eps_f(c, R), "epsg": f"={c}{R['eps']}/{pc}{R['eps']}-1"}
        for key, f in fs.items():
            ws[f"{c}{R[key]}"] = f; ws[f"{c}{R[key]}"].number_format = ws[f"G{R[key]}"].number_format; ws[f"{c}{R[key]}"].fill = FADEFILL

def _shift_rows(wb, sheet, start, n):
    """把 sheet 第 start 列起的內容整體下移 n 列，並同步改寫所有公式與定義名稱中指向這些列的參照（等同插入列）。"""
    ws = wb[sheet]
    cell_ = r"\$?[A-Z]{1,3}\$?\d+"
    pat = _re.compile(r"(?<![A-Za-z0-9_.!$:])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(" + cell_ + r")(?::(" + cell_ + r"))?(?![A-Za-z0-9_(!])")
    one = _re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")
    def fix(f, own):
        parts = f.split('"')
        for i in range(0, len(parts), 2):
            def rep(m):
                pre = m.group(1)
                tgt = (pre is None and own) or (pre is not None and pre.strip("'!") == sheet)
                def sh(x):
                    a, rr = one.fullmatch(x).groups()
                    rr = int(rr)
                    return f"{a}{rr + n if (tgt and rr >= start) else rr}"
                return (pre or "") + sh(m.group(2)) + (":" + sh(m.group(3)) if m.group(3) else "")
            parts[i] = pat.sub(rep, parts[i])
        return '"'.join(parts)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = fix(c.value, w.title == sheet)
    for nm in list(wb.defined_names.keys()):
        d = wb.defined_names[nm]
        if sheet + "!" in d.attr_text: d.attr_text = fix(d.attr_text, False)
    cells = [(c.row, c.column) for row in ws.iter_rows(min_row=start) for c in row if c.value is not None or c.has_style]
    for r, col in sorted(cells, reverse=True):
        src = ws.cell(r, col); dst = ws.cell(r + n, col)
        dst.value = src.value; dst._style = _copy(src._style); dst.number_format = src.number_format
        src.value = None; src._style = _copy(ws.cell(1, 30)._style)

def _tp_section(V, eps_ref, fwd, fwd_note, peers, peer_note, cons_note, method, pe_rng, fmt, cons_feps_note):
    """Val 第 0 節（第 4–48 列）：0.1 DCF、0.2 Comparable、0.3 結論。列標籤依工作單逐字；B／C／D＝情境 A 順利／B 平平／C 不如預期。"""
    T = {}
    def tset(r, key, label, fs, nf, note=""):
        V.cell(r, 1, label); T[key] = r
        for c, f in zip("BCD", fs):
            if f is None: continue
            V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = nf
        if note: V[f"E{r}"] = note
    V.cell(4, 1, "0. 目標價").font = BOLD
    V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
    hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
    V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法（PL 頁灰底）"
    for i, y in enumerate(YRS6):
        tset(8 + i, f"eps{y}", f"  FY{y}", [f"={eps_ref(s, y)}" for s in SC], "0.00")
    V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）")
    tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=IF(tp_roe<=tp_g,0.8,MIN(tp_g/tp_roe,0.8))", None, None], "0.0%", "ROE＝In 頁 tp_roe；tp_roe ≤ tp_g 時取上限 0.8")
    for i, y in enumerate(YRS6):
        tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{8 + i}*(1-$B$15)" for c in "BCD"], "0.00")
    tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], fmt,
         "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027")
    tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], fmt)
    tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], fmt, "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
    V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
    hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
    assert 3 <= len(peers) <= 5, peers
    for i, (nm, tk, pe, dt, src) in enumerate(peers, 28):
        V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
        V.cell(i, 4, dt); V.cell(i, 5, src); V.cell(i, 6, "Interested-party")
    tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{27 + len(peers)})", None, None], "0.0", peer_note)
    tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）",
         ["=" + "+".join(f"tp_w{s}*(" + "+".join(f"{w}*{c}{8 + YRS6.index(y)}" for y, w in fwd) + ")" for s, c in zip(SC, "BCD")), None, None], "0.00", fwd_note)
    tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], fmt, "同業中位數 × 前瞻 EPS")
    V.cell(37, 1, "0.3 結論").font = BOLD
    hdr(V, 38, ["項目", "數值", "", "", "說明"])
    tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], fmt, "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
    tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], fmt)
    tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], fmt, cons_note)
    tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px")
    tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
    tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", cons_feps_note)
    tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
    V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
    tset(46, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D23,"{fmt}")&"～"&TEXT(B23,"{fmt}")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
    tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
         [f"=IF(px/{c}9>MAX({pe_rng}),1,IF(px/{c}9<MIN({pe_rng}),0,PERCENTRANK({pe_rng},px/{c}9)))" for c in "BCD"], "0%",
         "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 6 節採用來源）的分位；高於十年最大值＝100%；只作檢查，不進目標價")
    V.cell(48, 1, "目標價方法"); V["B48"] = method
    for r_ in (24, 35, 39):
        V[f"B{r_}"].fill = KEY
    return T

def _tp_sens(SE, T):
    """Sens：目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.3／0.7）；放在案例摘要與第 40 列推導區塊之間。"""
    last = max(r for r in range(1, 40) if any(SE.cell(r, j).value is not None for j in range(1, 13)))
    t0 = last + 2
    SE.cell(t0, 1, "目標價參數（只改 DCF 參數與權重；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
    hdr(SE, t0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
    def dcf_f(c, r, g, reinv):
        cf = lambda i: f"Val!{c}{8 + i}*(1-{reinv})"
        pv = "+".join(f"{cf(i)}/(1+{r})^{i + 1}" for i in range(6))
        return f"({pv}+{cf(5)}*(1+{g})/({r}-{g})/(1+{r})^6)"
    cases = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
             ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
             ("DCF 權重 0.3", "Comparable 0.7", dict(w="0.3")), ("DCF 權重 0.7", "Comparable 0.3", dict(w="0.7"))]
    for k, (nm, st, kw) in enumerate(cases):
        i = t0 + 2 + k; r = kw.get("r", "tp_r"); g = kw.get("g", "tp_g"); w = kw.get("w", "tp_w_dcf")
        reinv = f"IF(tp_roe<={g},0.8,MIN({g}/tp_roe,0.8))"
        SE.cell(i, 1, nm); SE.cell(i, 2, st)
        SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_f(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
        SE[f"D{i}"] = f"=Val!B{T['comptp']}"
        SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
        SE[f"F{i}"] = f"=E{i}/E${t0 + 2}-1"
        for c in "CDE": SE[f"{c}{i}"].number_format = "0.00"
        SE[f"F{i}"].number_format = "0.0%"
    assert t0 + 1 + len(cases) < 40, t0
    return t0, t0 + 1 + len(cases)

def _fade_all(eps_f, nim10):
    """PL：三情境推導區塊與頂端摘要延伸到 FY2032；In 頁 fade_m 連到 PL 頁 FY2029 淨利率。nim10＝十年平均淨利率的名稱；None＝未找到，依工作單用 FY2029 × 0.8。"""
    _move_notes(PL)
    for s in SC:
        _fade(PL, PB[s], s, eps_f)
        for key in ("rev", "g", "nim", "ni", "eps", "epsg"):
            if (s, key) not in SUMR: continue
            for y in FADE:
                cc = PL[f"{FCX[y]}{SUMR[(s, key)]}"]; cc.value = f"={FCX[y]}{PB[s][key]}"
                cc.number_format = PL[f"G{SUMR[(s, key)]}"].number_format; cc.fill = FADEFILL
    nimr = {s: f"PL!$G${PB[s]['nim']}" for s in SC}
    IN.cell(INR["fade_mA"], 5, f"={nimr['A']}")
    IN.cell(INR["fade_mB"], 5, f"=({nimr['B']}+{nim10})/2" if nim10 else f"=({nimr['B']}+{nimr['B']}*0.8)/2")
    IN.cell(INR["fade_mC"], 5, f"={nim10}" if nim10 else f"={nimr['C']}*0.8")
    PL["A2"] = ("灰底＝FY2030–FY2032 收斂法，非驅動因子推導（上游市場層只到 2029）：營收年增由 FY2029 線性收斂到 fade_g（A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt），"
                "淨利率由 FY2029 線性收斂到 fade_m（In 頁 K 節），FY2032 達到；股數依 sh_g")

def _rule_pe_range(V):
    """第 6 節十年倍數表（2016–2025 列）中採用來源的欄位範圍（PERCENTRANK 用）。"""
    yr = [r for r in range(1, V.max_row + 1) if V.cell(r, 1).value in [str(y) for y in range(2016, 2026)]]
    yr = [r for r in yr if r > 50][:10]
    assert len(yr) == 10 and yr[-1] - yr[0] == 9, yr
    return f"$B${yr[0]}:$B${yr[-1]}"

_fade_all(lambda c, R: f"{c}{R['ni']}/{c}{R['sh_']}", None)
_shift_rows(wb, "Val", 4, 47)
TPT = _tp_section(V, lambda s_, y: f"PL!{FCX[y]}{SUMR[(s_, 'eps')]}", [(2027, "2/12"), (2028, "10/12")],
    "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）",
    [("新易盛（Eoptolink）", "300502.SZ", 12.52, "2026-10-23", "stockanalysis forward PE（S&P Global 共識；價格日期如左）"),
     ("天孚通信（TFC）", "300394.SZ", 46.53, "2026-09-30", "同上（頁面價格日期較舊）"),
     ("光迅科技（Accelink）", "002281.SZ", 77.09, "2026-08-18", "同上（頁面價格日期較舊）"),
     ("華工科技（HGTECH）", "000988.SZ", 45.08, "2026-08-07", "同上（頁面價格日期較舊）"),
     ("劍橋科技（CIG）", "603083.SS", 31.74, "2026-10-23", "同上")],
    "A 股光通訊同業（中際旭創是 A 股公司，同業只用 A 股；工作單 Andy 11:53）",
    "stockanalysis 預測頁：14 位平均 1,321、中位數 1,400（2026-10-25 擷取）；多數訂於 10 月股價下跌前",
    ("DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為近三年平均約 31%（再投資率約 13%），FY2027–32 明確期＋永續，FY2030 起為收斂法。"
     "Comparable：A 股光通訊同業（新易盛、天孚通信、光迅科技、華工科技、劍橋科技）前瞻 P/E 中位數；同業規模較小、部分價格日期為 8–9 月，"
     "而新易盛只有約 13 倍，中位數偏高估。情境對應：顆數情境 A＝順利、B＝平平、C＝不如預期；FY2027 三情境 EPS 只差在網通 CPO 情境（A＝CPO 替代快），顆數差異從 CY2028 起。"),
    _rule_pe_range(V), "#,##0.00",
    "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝In 頁 cons_feps（同花順 FY2027、FY2028 共識，與本模型同期間加權）")
TPS0, TPS_LAST = _tp_sens(SE, TPT)

# ---------------- v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 14:00、14:17；既有數值與儲存格位置不變） ----------------
_UPW = [r_ for r_ in range(1, 60) if V.cell(r_, 1).value == "上檔空間：加權目標價"]
assert _UPW == [42], _UPW
GR_ROW, RC_ROW = 49, 50   # 0.3 結論後的空白預留列（不移動其他列）
for _r in (GR_ROW, RC_ROW):
    assert all(V.cell(_r, _c).value is None for _c in range(1, V.max_column + 1)), _r
V.cell(GR_ROW, 1, "上檔空間級距")
V.cell(GR_ROW, 2, '=IF(B42>0.5,"大於 50%",IF(B42>0.25,"25%–50%",IF(B42>=0,"0%–25%",IF(B42>=-0.25,"0%～−25%",IF(B42>=-0.5,"−25%～−50%","小於 −50%")))))').number_format = "@"
V.cell(GR_ROW, 5, "以「上檔空間：加權目標價」分六級（Andy 2026-10-25 14:00）")
V.cell(RC_ROW, 1, "投資建議")
V.cell(RC_ROW, 2, '=IF(B42>0.5,"強力買進",IF(B42>0.25,"買進",IF(B42>=-0.25,"中立",IF(B42>=-0.5,"放空","強力放空"))))').number_format = "@"
V.cell(RC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
# 自由現金流（算力收支 2026-10-25 請求；地圖抽數依列標籤讀取）：In 頁、PL 頁、Sources 頁皆在現有最後一列之後新增，不移動既有列
_ri = IN.max_row + 1
for _k, _lab, _f, _s1 in (("fcf_da", "折舊 ÷ 營收（自由現金流用；各年營收 × 比率）", "=819.47/38240", "FY2025 年報：折舊與攤銷 819.47 百萬元（stockanalysis 轉載現金流量表）÷ 營收 38,240 百萬元（Verified）"),
                         ("fcf_capex", "資本支出 ÷ 營收（自由現金流用；各年營收 × 比率）", "=2759.99/38240", "FY2025 年報：購建固定資產、無形資產和其他長期資產支付的現金 2,759.99 百萬元（中財網；stockanalysis 2,760）÷ 營收 38,240 百萬元（Verified）")):
    IN.cell(_ri, 1, _k); IN.cell(_ri, 2, _lab); IN.cell(_ri, 3, "%"); IN.cell(_ri, 5, _f).number_format = "0.00%"
    IN.cell(_ri, 7, "Assumed"); IN.cell(_ri, 8, _s1); IN.cell(_ri, 9, "FY2026 起沿用最近一個完整年度（FY2025）年報比率（Assumed）；只用於 PL 頁自由現金流列，不進 EPS 與目標價")
    defname(_k, f"In!$E${_ri}"); _ri += 1
_fr = PL.max_row + 2
PL.cell(_fr, 1, "自由現金流（v0.2.1；算力收支 2026-10-25 請求）：歸母淨利＋折舊 − 資本支出；折舊、資本支出＝各年營收 × In 頁 fcf_da、fcf_capex（FY2025 年報比率）；不含營運資金變動、不進 EPS 與目標價").font = BOLD
_yrs = sorted(FCX)
hdr(PL, _fr + 1, ["項目", "單位"] + [f"FY{y}" + ("A" if y == 2025 else "") for y in _yrs] + ["說明"])
for _i, s in enumerate(SC):
    _rr = _fr + 2 + _i
    PL.cell(_rr, 1, f"自由現金流（淨利＋折舊 − 資本支出；不含營運資金變動）：情境 {s}"); PL.cell(_rr, 2, "RMB B")
    for y in _yrs:
        PL[f"{FCX[y]}{_rr}"] = f"={FCX[y]}{SUMR[(s, 'ni')]}+{FCX[y]}{SUMR[(s, 'rev')]}*(fcf_da-fcf_capex)"; PL[f"{FCX[y]}{_rr}"].number_format = "#,##0.00"
    PL.cell(_rr, 3 + len(_yrs), "FY2025A 為實績營收 × 年報比率；FY2030 起營收、淨利為收斂法" if _i == 0 else None)
SO.cell(SO.max_row + 1, 1, "[Verified] 公司 FY2025 年報現金流量表（資料商轉載）：折舊與攤銷 819.47、購建長期資產現金 2,759.99（百萬元）；營收 382.4 億元——自由現金流比率錨點"); SO.cell(SO.max_row, 2, "https://quote.cfi.cn/xjll/14097/300308.html（現金流量表 2025 年報：購建固定資產、無形資產和其他長期資產支付的現金 275,999.47 萬元）；https://stockanalysis.com/quote/she/300308/financials/cash-flow-statement/（FY2025 折舊與攤銷 819.47 百萬元、資本支出 2,760 百萬元）；https://stock.10jqka.com.cn/20260331/c675646713.shtml（2025 年營收 382.4 億元）")
SO.cell(SO.max_row + 1, 1, "[Assumed] 自由現金流：FY2026 起折舊 ÷ 營收、資本支出 ÷ 營收沿用 FY2025 比率（In 頁 fcf_da、fcf_capex）；不含營運資金變動"); SO.cell(SO.max_row, 2, "本模型假設")
OBS = [
       "目標價 ¥1,343（DCF 823、Comparable 1,864 各半），比現價 612.63（2026-10-23）高 119.3%，級距 大於 50%；Comparable 單獨為 +204.2%。主因是倍數不是 EPS：本模型 12 個月後前瞻 EPS 41.34 比共識前瞻 EPS 75.09 低 44.9%，但 A 股光通訊同業（新易盛、天孚、光迅、華工、劍橋）前瞻 P/E 中位數 45.1 倍，現價只隱含 12.1 倍（共識未來 12 月 EPS）、共識目標價 1,321 隱含 17.6 倍；Comparable 未扣同業規模較小、部分價格日期較舊的溢價，也未計美國限制風險。DCF 單獨也有 +34.3%（ROE 30.6% 高於折現率 10%）。",
       "最敏感的是美國限制與份額：美國限制案例（份額相對值 2028 0.55）FY2028 EPS −38.2%、份額流失 −21.2%、降價快於模型 −19.0%、AI 市場 ×0.8 −18.6%；情境 C FY2028 EPS 29.54，比情境 A 的 48.26 低 38.8%，DCF 價格帶 504～1,017。",
       "追蹤 3Q26 季報（10-31 前）：FY2026 錨點的 2H26（季增 15%、毛利率 47%）為假設。網通 CPO 已改三情境（v0.2.2；A＝替代快、C＝慢）：CPO 快對中際旭創是減分，情境 A FY2028 EPS 由 49.78 降到 48.26（−3.1%）、C 由 29.17 升到 29.54（+1.2%）、B 不變；CPO 份額 0 時 FY2028 EPS −4.1%。"
      ]   # 重點觀察（與 README.md 相同文字）
_n0 = RD.max_row
RD.cell(_n0 + 1, 1, "v0.2.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（Val 頁 0.3 結論加「上檔空間級距」「投資建議」兩列、README 頁加重點觀察三點；數值不變，既有儲存格位置不變）。另在 PL 頁末新增「自由現金流（淨利＋折舊 − 資本支出；不含營運資金變動）」三情境列（FY2025A–FY2032；歸母淨利口徑；算力收支 2026-10-25 請求），In 頁末新增 fcf_da、fcf_capex（FY2025 年報比率 2.14%、7.22%）。")
RD.cell(_n0 + 2, 1, f'="投資建議："&Val!B{RC_ROW}&"（依上檔空間固定規則）"').font = BOLD
for _i, _t in enumerate(OBS, 1): RD.cell(_n0 + 2 + _i, 1, f"重點觀察 {_i}：" + _t)

order =["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
# ---------------- v0.2.2（2026-10-26）：網通 CPO 三情境（同字母）；In 頁在現有最後一列之後新增，不移動既有列 ----------------
_ri = IN.max_row + 2
IN.cell(_ri, 1, "網通 CPO 三情境（v0.2.2；markets/network v0.2.3：A＝替代快、B＝平平、C＝慢；情境 A 為上方 A 節，B、C 為下列，各再乘同字母情境顆數係數）").font = BOLD; _ri += 1
for _kb, _lab, _kind in [('mpl', 'AI 可插拔光模組市場', 'mpl'), ('mcp', 'AI CPO 光學市場', 'mcp')]:
    for _s in ("B", "C"):
        for _j, _y in enumerate((26, 27, 28, 29)):
            _k = f"{_kb}{_s}_{_y}"
            IN.cell(_ri, 1, _k); IN.cell(_ri, 2, f"情境 {_s}：{_lab} CY20{_y}"); IN.cell(_ri, 3, "US$B")
            IN.cell(_ri, 5, NETS[(_kind, _s)][0][_j]).font = BLUE
            IN.cell(_ri, 7, "Derived"); IN.cell(_ri, 8, NETS[(_kind, _s)][1])
            IN.cell(_ri, 9, f"網通模型 CPO 情境 {_s}（技術替代屬行業情境，公司讀同字母）" if _y == 26 else "同上")
            wb.defined_names[_k] = DefinedName(_k, attr_text=f"In!$E${_ri}"); _ri += 1
_n1 = RD.max_row
RD.cell(_n1 + 1, 1, 'v0.2.2（2026-10-26）：網通模型 v0.2.3 的 CPO 滲透改三情境（A＝替代快、B＝平平、C＝慢），本模型 A／B／C 改讀同字母的可插拔與 CPO 市場（In 頁末新增 B、C 兩組；情境 A 沿用 A 節 mpl_／mcp_，數值改為網通情境 A）。CPO 快對中際旭創是減分（可插拔市場較小、CPO 份額只部分抵銷），所以 A 情境 EPS 降幅最大；FY2027 起三情境不再相同。其他輸入不變。')
wb.save(sys.argv[1])
if os.environ.get("INNOLIGHT_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK_ROW, "SENSN": SENS_N, "IMP": IMP, "VAL": VAL, "IMPL": {str(k): v for k, v in IMPL.items()}, "SB": SB},
              open(os.environ["INNOLIGHT_ROWMAP"], "w"), ensure_ascii=False, indent=1)
