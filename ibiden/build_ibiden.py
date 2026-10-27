# Ibiden（4062）收支與估值模型 v0.2.1（AI 半導體 Project：載板與 PCB 次產業樣板）
# 用法：python3 ibiden/build_ibiden.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度 4 月至隔年 3 月：FY(t)＝t/4–(t+1)/3（公司稱「(t+1) 年 3 月期」），對照 CY(t) 的加速器顆數（重疊 9 個月，不做平移）。
# 金額：十億日圓（¥B）；EPS：日圓（2026-10-01 一拆二之後的股數口徑）。美元換算只作對照（¥152／US$，公司 FY2026 前提）。
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
UNITS = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith("u_") or str(k).startswith("g29_")):
        UNITS[k] = (r[4].value, r[1].value, f"{rel(NET_F)}｜In!E{r[0].row}")
NETR = {}
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    if r[0].value == "AI IT 資本支出（加速器 × 每顆 IT capex）":
        NETR["itc"] = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert "itc" in NETR
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
NV = ("bw", "rb", "ru", "hp"); OT = ("amd", "tpu", "trn", "oth")
assert all(f"u_{g}_{y}" in UNITS for g in NV + OT for y in (26, 27, 28)), sorted(UNITS)

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億日圓（¥B）；顆數：百萬顆；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝Ibiden 決算短信與補足資料（TDnet）；公司預測為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
TAN25 = "公司：2026 年 3 月期決算短信（Sources 頁）"
TAN25_URL = "https://www.ibiden.com/ir/items/ConsolidatedFinancialResultsFY2025.pdf"
TANQ1 = "公司：2027 年 3 月期 Q1 決算短信（2026-08-04；Sources 頁）"
TANQ1_URL = "https://finance.biggo.jp/news/jpx_tdnet_140120260731504718"
SUPQ1 = "公司：Q1 補足資料（分部門預測、設備投資、匯率前提；Sources 頁）"
SUPQ1_URL = "https://finance.biggo.jp/news/jpx_tdnet_140120260803507041"
STD = "SEMICON.TODAY 2026-08"
LIMO = "LIMO 2026-08-15"
TOMS = "Tom's Hardware：The state of ABF substrates in 2026"
DG = "Data Gravity 2026-06-14"

sec("A. 加速器顆數（讀 markets/network In 頁；CY2026–CY2028 為輸入、CY2029＝CY2028 ×（1＋成長率））")
for g in NV + OT:
    for y in (26, 27, 28):
        v, lab, src = UNITS[f"u_{g}_{y}"]
        inp(f"u_{g}_{y}", lab, "百萬顆", None, v, None, "Derived", src, "唯讀上游；不在本模型另訂")
    v, lab, src = UNITS[f"g29_{g}"]
    inp(f"g29_{g}", lab, "%", None, v, None, "Derived", src, "同上")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, NETR["itc"][0][j], None, "Derived", NETR["itc"][1], "Recon 用：Ibiden AI 營收 ÷ AI IT 資本支出（每 GW 對照）" if y == 26 else "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、NVIDIA、Broadcom、Arista 模型相同）")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")

sec("B. Ibiden 份額與每顆載板金額（本專案權威：晶片商載板份額，規則第 13 條）")
inp("w_ot", *SIL.row9("abf_wot_ib")[1:])
inp("ai_e", "FY2026 AI 伺服器載板 ÷ 電子事業營收", "%", 0.55, 0.65, 0.75, "Assumed",
    "公司未揭露：未找到（決算短信只說「生成 AI 伺服器向け・汎用サーバー向け高機能 IC パッケージ基板」帶動）", "semi-connect 2026：AI 高階載板營收多數為 NVIDIA 向（估計，Interested-party）")
inp("g_asp", *SIL.row9("abf_gasp")[1:])
inp("rel_27", *SIL.row9("abf_ibrel_27")[1:])
inp("rel_28", *SIL.row9("abf_ibrel_28")[1:])
inp("rel_29", *SIL.row9("abf_ibrel_29")[1:])

sec("C. 短缺溢價（價格指數；2026＝1）與 AI 載板利潤率")
inp("pp_27", *SIL.row9("abf_pp_27")[1:])
inp("pp_28", *SIL.row9("abf_pp_28")[1:])
inp("pp_29", *SIL.row9("abf_pp_29")[1:])
inp("m_ne", "電子事業非 AI 部分營業利益率", "%", 0.10, 0.15, 0.20, "Assumed", "FY2024 電子事業營益率 13.6%（26.8 ÷ 197.2，AI 比重較低時；" + TAN25 + "）", "FY2025 18.6%（45.2 ÷ 243.3）")
inp("g_ne", "電子事業非 AI 營收年增（通用伺服器、PC 載板等）", "%", 0.00, 0.05, 0.10, "Assumed", STD + "：通用伺服器市場緩步成長", "PC：未找到")

sec("D. 陶瓷、其他事業與全公司")
inp("g_cer", "陶瓷事業營收年增（DPF、汽車排氣）", "%", -0.05, 0.00, 0.03, "Assumed", "FY2025 −1.8%（84.1 → 82.6；" + TAN25 + "）", "FY2026 預測 84.0（+1.8%；" + SUPQ1 + "）")
inp("m_cer", "陶瓷事業營益率", "%", 0.05, 0.077, 0.10, "Derived", "FY2026 預測 6.5 ÷ 84.0（" + SUPQ1 + "）", "FY2025 9.3%（7.6 ÷ 82.6）")
inp("g_oth", "其他事業營收年增（建材等）", "%", 0.00, 0.03, 0.05, "Assumed", "FY2025 +2.5%（88.1 → 90.3）", "FY2026 預測 91.0（+0.7%）")
inp("m_oth", "其他事業營益率", "%", 0.09, 0.115, 0.13, "Derived", "FY2026 預測 10.5 ÷ 91.0（" + SUPQ1 + "）", "FY2025 9.9%（9.0 ÷ 90.3）")
inp("ni_oi", "淨利 ÷ 營業利益（稅、非營業、少數股權）", "x", 0.62, 0.661, 0.70, "Derived", "FY2026 預測 84 ÷ 127（" + SUPQ1 + "）", "FY2025 1.03（含政策持股出售利益，不具代表性）")
inp("sh_g", "股數年變動（2031 年到期 CB 轉換）", "%", 0.00, 0.00, 0.03, "Assumed", "CB 轉換價 ¥2,243.1（分割後；" + STD + "），遠低於現價", "CB 發行額：未找到；FY2025 稀釋 EPS 比基本低 5.8%（214.91 ÷ 228.16）")
inp("fx", "匯率（美元換算用）", "¥／US$", None, 152, None, "Interested-party", SUPQ1 + "：FY2026 前提 ¥152", "—")

sec("E. FY2026 錨點（公司 2026-08-04 上修後預測；Q2 決算後更新）")
inp("f_rev", "FY2026 營收預測", "¥B", None, 550.0, None, "Interested-party", SUPQ1, LIMO + "：5,500 億（+32.1%）")
inp("f_e", "FY2026 電子事業營收預測", "¥B", None, 375.0, None, "Interested-party", SUPQ1, "Q1 實績 77.5（年增 37.7%；" + STD + "）")
inp("f_eoi", "FY2026 電子事業營業利益預測", "¥B", None, 110.0, None, "Interested-party", SUPQ1, "Q1 實績 21.4（營益率 27.6%）")
inp("f_cer", "FY2026 陶瓷事業營收預測", "¥B", None, 84.0, None, "Interested-party", SUPQ1, "—")
inp("f_ceroi", "FY2026 陶瓷事業營業利益預測", "¥B", None, 6.5, None, "Interested-party", SUPQ1, "—")
inp("f_oth", "FY2026 其他事業營收預測", "¥B", None, 91.0, None, "Interested-party", SUPQ1, "—")
inp("f_othoi", "FY2026 其他事業營業利益預測", "¥B", None, 10.5, None, "Interested-party", SUPQ1, "—")
inp("f_oi", "FY2026 營業利益預測", "¥B", None, 127.0, None, "Interested-party", SUPQ1, STD + "：900 → 1,270 億")
inp("f_ni", "FY2026 淨利預測（歸屬母公司）", "¥B", None, 84.0, None, "Interested-party", SUPQ1, STD + "：580 → 840 億")
inp("f_eps", "FY2026 EPS 預測（分割後）", "¥", None, 149.68, None, "Interested-party", STD + "：反映 2026-10-01 一拆二", "—")

sec("F. FY2025 實績（2026 年 3 月期）")
for k, lab, v in (("a_rev", "營收", 416.201), ("a_e", "電子事業營收", 243.316), ("a_eoi", "電子事業營業利益", 45.248), ("a_cer", "陶瓷事業營收", 82.554), ("a_ceroi", "陶瓷事業營業利益", 7.646),
                  ("a_oth", "其他事業營收", 90.330), ("a_othoi", "其他事業營業利益", 8.964), ("a_oi", "營業利益", 62.027), ("a_gp", "毛利", 131.668), ("a_ni", "淨利（歸屬母公司）", 63.713)):
    inp(k, f"FY2025 {lab}", "¥B", None, v, None, "Verified", TAN25, "SEMICON.TODAY、semi-connect 同值" if k in ("a_rev", "a_e", "a_oi") else "—")
inp("a_eps", "FY2025 基本 EPS（分割後換算＝228.16 ÷ 2）", "¥", None, "=228.16/2", None, "Derived", TAN25 + "：228.16（2026-01 一拆二後、2026-10 一拆二前）", "2026-10-01 再一拆二")

sec("G. 價格與市場資料（只作 Val、Recon 對照）")
inp("px", "股價（2026-10-23 收盤，分割後）", "¥", None, 8512, None, "Interested-party", "markets/invest_map/data/prices.csv：株探、松井證券（兩來源核對；v0.2 更新）", "v0.1.1 為 10-16 收盤 9,764（stockanalysis）；前瞻 P/E 隱含 EPS 仍以 9,764 計（In 頁 cons_feps）")
inp("pe_fwd_sa", "前瞻 P/E（stockanalysis）", "x", None, 53.97, None, "Interested-party", "stockanalysis（口徑：未來 12 個月共識，頁面未明示）", "第二來源：未找到")
inp("pe_ttm", "近 12 月 P/E", "x", None, 83.85, None, "Interested-party", "stockanalysis", "—")
inp("pe_fair_andy", "Andy 先驗：合理 P/E", "x", None, 25, None, "Assumed", "先驗_Andy_Ibiden_20261024 I10b（Drive；不在本 repo）", "只作對照")
inp("tgt30_rev", "公司 FY2030 目標營收", "¥B", None, 1000, None, "Interested-party", "Quartr FY2025 摘要（AI 摘要；營收數字有誤處未採用）", "第二來源：未找到")
inp("tgt30_oi", "公司 FY2030 目標營業利益", "¥B", None, 300, None, "Interested-party", "同上", "第二來源：未找到")
sec("H. 目標價情境權重（v0.1.1 新增；Val 頁第 0 節用）")
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（{dict(A='順利', B='平平', C='不如預期')[s]}）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%；v0.2 起用於 DCF 與 Comparable 的情境加權（Val 第 0 節）")

# ---------------- v0.2 設定與新增輸入（附加在 In 既有輸入之後，不位移既有列） ----------------
V2 = {'name': 'Ibiden',
 'unit': '¥B',
 'shunit': '十億股',
 'cur': '¥',
 'shg': 'sh_g',
 'epsfmt': '#,##0.0',
 'pxfmt': '#,##0',
 'keys': {'rev': 'rev', 'ni': 'ni', 'sh': 'sh', 'eps': 'eps', 'sh25': False},
 'roe': ('=AVERAGE(0.1217,0.0682,0.0685)',
         'stockanalysis 比率頁（Return on Equity）：2026 年 3 月期 12.17%、2025 年 3 月期 6.82%、2024 年 3 月期 6.85%（即本模型 FY2025、FY2024、FY2023）',
         '單一來源；近 12 月 12.65%（同站，2026-08-27）'),
 'cons_tp': (12432.35, 'stockanalysis 個股頁（S&P Global 彙整）18 位分析師平均目標價 12,432（頁面價格日期 2026-10-16；分割後）', '第二來源：未找到'),
 'cons_feps': ('=9764/53.97', 'stockanalysis 個股頁 2026-10-16：價格 9,764 ÷ Forward PE 53.97（即 In 頁 pe_fwd_sa）'),
 'peers': [('欣興（Unimicron）', '3037.TW', 42.28, '2026-10-22'),
           ('南電（Nan Ya PCB）', '8046.TW', 33.73, '2026-09-07'),
           ('景碩（Kinsus）', '3189.TW', 44.7, '2026-10-22'),
           ('三星電機（Samsung Electro-Mechanics）', '009150.KS', 43.4, '2026-08-28'),
           ('AT&S', 'ATS.VI', 14.91, '2026-10-23')],
 'peer_note': 'ABF／FC-BGA 載板同業：欣興、南電、景碩、三星電機、AT&S（券商常用的載板比較組；新光電氣 2025 年下市未列）。非 A 股公司的同業名單不放中國 A 股（A 股有結構性溢價；Andy 2026-10-25 11:53）',
 'fwd_w': (0.5, 0.5),
 'fwd_note': '會計年度 4 月至隔年 3 月：約 2027-10 時的未來 12 月（2027-10～2028-09）≈ 6／12 × FY2027＋6／12 × FY2028',
 'pe10': ([(2016, 44.1), (2017, None), (2018, 35.3), (2019, 137.0), (2020, 55.2), (2021, 53.1), (2022, 40.7), (2023, 26.8), (2024, 57.2), (2025, 33.6)],
          'companiesmarketcap 日曆年年底 P/E，Interested-party；2017 年虧損不計'),
 'px_note': '現價＝In 頁 px（2026-10-23 收盤 8,512，分割後）',
 'method': 'DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE＝近三年平均約 8.6%（再投資率約 46%），FY2027–29 驅動因子、FY2030–32 收斂法＋永續；基準日 2026 年底（FY2027 實際於 2028-03 結束，未做半年調整）。Ibiden 資本密集（FY2026 設備投資 2,100 '
           '億），但模型無折舊與資本支出明細，未改用淨利＋折舊－資本支出，現金流偏高估。Comparable：欣興、南電、景碩、三星電機、AT&S 前瞻 P/E 中位數；台系載板股含 AI 題材溢價（約 34–45x），部分價格為 8–9 月，偏高估。情境對應：A＝順利、B＝平平、C＝不如預期（原顆數情境）。',
 'readme_val': '估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − tp_g ÷ tp_roe），FY2027–FY2032 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，推到 12 個月後）＋0.5 × Comparable 目標價（欣興、南電、景碩、三星電機、AT&S 前瞻 P/E 中位數 × 12 '
               '個月後前瞻 EPS）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解（EPS 因素、倍數因素）、DCF 價格帶、十年 P/E 分位檢查（只作檢查）。會計年度 4 月制：12 個月後前瞻 EPS＝½ FY2027＋½ FY2028。',
 'readme_add': ['v0.2（2026-10-25）：(1) PL 頁最下方新增延伸表，三情境延伸到 FY2032；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 A＝tp_g＋2pt、B＝tp_g、C＝tp_g−2pt；淨利率由 FY2029 收斂到 A＝FY2029、B＝FY2029 '
                '與十年平均的中點、C＝十年平均；十年平均淨利率未找到，依工作單以 FY2029 × 0.8 替代，Assumed；PL 延伸表灰底）。(2) Val 第 0 節目標價＝DCF 與 Comparable 加權（tp_w_dcf＝0.5）；v0.1.1 的「0. 目標價與上檔空間」改名「舊法：規則值情境加權（參考）」保留，不再是目標價。(3) Sens '
                '新增目標價參數（tp_r ±1pt、tp_g ±1pt、tp_w_dcf 0.4／0.8）。(4) Recon 新增目標價對帳。(5) 現價更新為 2026-10-23 收盤 ¥8,512（markets/invest_map/data/prices.csv：株探、松井證券兩來源）；前瞻 P/E 隱含 EPS 改以原取得價格 9,764 計（In 頁 '
                'cons_feps），數值不變。 新增內容一律放在各頁既有內容之後，FY2025–FY2029 既有數值不變。',
                '情境對應：A＝順利（顆數基準）、B＝平平（CY2028 顆數 ×0.8、CY2029 ×0.85）、C＝不如預期（CY2028–29 ×0.6），即原本的三個顆數情境。三情境在 FY2028 前完全相同（只差顆數係數，從 CY2028 起），DCF 三情境仍照算，差異來自 FY2028 之後（含 FY2030 起的收斂法）。',
                '先驗：v0.1 的 Andy 合理 P/E 25x 對照（Val 第 4 節）保留不動，不進目標價；本版未新增先驗。']}
sec("J. 目標價（DCF 與 Comparable 加權）與收斂法參數（v0.2 新增；黃底＝Assumed，Andy 可改）")
inp("tp_r", "折現率／要求報酬（DCF）", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 10%；美、台、日、韓同一值）", "Sens 頁 ±1pt")
inp("tp_g", "長期成長率（DCF 永續；情境 B 的 FY2032 營收年增收斂目標）", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "Sens 頁 ±1pt（只動 DCF，PL 收斂目標不隨之變動）")
inp("tp_roe", "ROE（近三年平均）", "%", None, V2["roe"][0], None, "Derived", V2["roe"][1], V2["roe"][2])
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 改為 0.5：DCF／Comparable 各半）", "Sens 頁 0.4／0.8")
inp("fade_gA", "情境 A（順利）：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂，FY2032 達到")
inp("fade_gB", "情境 B（平平）：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C（不如預期）：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("nim10_f", "十年平均淨利率的替代：最後一年（FY2029）淨利率 ×（本值）", "x", None, 0.8, None, "Assumed", "十年（2016–2025）平均淨利率：未找到（免費資料商只到近五年）；依工作單以最後一年 × 0.8 替代", "淨利率收斂目標：A＝FY2029、B＝FY2029 與替代值的中點、C＝替代值（PL 延伸表）")
inp("cons_tp", "市場共識目標價（賣方平均）", V2["cur"], None, V2["cons_tp"][0], None, "Interested-party", V2["cons_tp"][1], V2["cons_tp"][2])
inp("cons_feps", "共識前瞻 12 月 EPS（＝價格 ÷ 前瞻 P/E）", V2["cur"], None, V2["cons_feps"][0], None, "Interested-party", V2["cons_feps"][1], "賣方共識，只作對帳與差距拆解，不作校準")

r = 4
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    r += 1
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 80, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — FY2025 實績與 FY2026 錨點（¥B）。FY2026 錨點＝公司 2026-08-04 上修後的全年預測（Q1 公布時；分部門預測取自 Q1 補足資料）。Q2 決算（預計 10 月底或 11 月初）後更新"
Q["A2"] = "公司未揭露 AI 伺服器載板營收；本模型以 In 頁 ai_e（Assumed）切分電子事業，並據此校準每顆有效金額與 AI 載板營業利益率（Derived）"
hdr(Q, 3, ["項目", "FY2025A", "FY2026 錨點", "年增", "標記", "說明"])
QR = {}
for i, (k, lab, a, f, fmt) in enumerate((("rev", "營收", "=a_rev", "=f_rev", "#,##0.0"), ("e", "  電子事業", "=a_e", "=f_e", "#,##0.0"), ("cer", "  陶瓷事業", "=a_cer", "=f_cer", "#,##0.0"),
                                          ("oth", "  其他事業", "=a_oth", "=f_oth", "#,##0.0"), ("oi", "營業利益", "=a_oi", "=f_oi", "#,##0.0"), ("eoi", "  電子事業", "=a_eoi", "=f_eoi", "#,##0.0"),
                                          ("ceroi", "  陶瓷事業", "=a_ceroi", "=f_ceroi", "#,##0.0"), ("othoi", "  其他事業", "=a_othoi", "=f_othoi", "#,##0.0"),
                                          ("adj", "  調整（全公司 − 三事業）", "=a_oi-a_eoi-a_ceroi-a_othoi", "=f_oi-f_eoi-f_ceroi-f_othoi", "#,##0.00"),
                                          ("ni", "淨利（歸屬母公司）", "=a_ni", "=f_ni", "#,##0.0"), ("eps", "EPS（分割後，日圓）", "=a_eps", "=f_eps", "#,##0.00")), 4):
    Q.cell(i, 1, lab); Q[f"B{i}"] = a; Q[f"C{i}"] = f; Q[f"D{i}"] = f"=C{i}/B{i}-1"
    for c in "BC": Q[f"{c}{i}"].number_format = fmt
    Q[f"D{i}"].number_format = "0.0%"
    Q.cell(i, 5, "Verified／Interested-party"); QR[k] = i
Q.cell(QR["ni"], 6, "FY2025 含政策持股出售利益（稅後約 345 億，Quartr 摘要），淨利高於營業利益")
r = QR["eps"] + 2
Q.cell(r, 1, "比率").font = BOLD; r += 1
for k, lab, fb, fc in (("em", "電子事業營益率", "=B{e}/B{s}".format(e=QR["eoi"], s=QR["e"]), "=C{e}/C{s}".format(e=QR["eoi"], s=QR["e"])),
                       ("om", "營業利益率", "=B{e}/B{s}".format(e=QR["oi"], s=QR["rev"]), "=C{e}/C{s}".format(e=QR["oi"], s=QR["rev"])),
                       ("esh", "電子事業 ÷ 營收", "=B{e}/B{s}".format(e=QR["e"], s=QR["rev"]), "=C{e}/C{s}".format(e=QR["e"], s=QR["rev"])),
                       ("gm", "毛利率（FY2026 預測未揭露）", "=a_gp/a_rev", None),
                       ("sh", "股數（十億股＝淨利 ÷ EPS）", f"=B{QR['ni']}/B{QR['eps']}", f"=C{QR['ni']}/C{QR['eps']}")):
    Q.cell(r, 1, lab); Q[f"B{r}"] = fb
    if fc: Q[f"C{r}"] = fc
    for c in "BC": Q[f"{c}{r}"].number_format = "0.0%" if k != "sh" else "0.0000"
    Q.cell(r, 5, "Derived"); QR[k] = r; r += 1
Q.cell(QR["sh"], 6, "FY2025 股數含政策持股出售前後變動，只作參考；預測期用 FY2026 錨點股數")
for nm, k in (("rev", "rev"), ("e", "e"), ("cer", "cer"), ("oth", "oth"), ("oi", "oi"), ("eoi", "eoi"), ("ni", "ni"), ("eps", "eps"), ("adj", "adj")):
    defname(f"q25_{nm}", f"Q!$B${QR[k]}"); defname(f"q26_{nm}", f"Q!$C${QR[k]}")
defname("q26_sh", f"Q!$C${QR['sh']}")
Q.column_dimensions["A"].width = 40
for c in "BCD": Q.column_dimensions[c].width = 14
Q.column_dimensions["E"].width = 26; Q.column_dimensions["F"].width = 80
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G"}
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
def units_f(g, y):
    return f"u_{g}_{y % 100}" if y <= 2028 else f"u_{g}_28*(1+g29_{g})"
DEFP = dict(scn="A", w="w_ot", aie="ai_e", ga="g_asp", rel="rel", pp="pp", mne="m_ne", gne="g_ne", gcer="g_cer", goth="g_oth", nioi="ni_oi", shg="sh_g")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 8, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 伺服器載板（D 欄＝FY2026 校準年；FY(t) 對照 CY(t) 顆數）")
    row("unv", "NVIDIA 加速器顆數 × 情境係數", "百萬顆", {c_(y): "=(" + "+".join(units_f(g, y) for g in NV) + f")*{umx(y, s)}" for y in BU}, note="markets/network In：Blackwell、Rubin、Rubin Ultra、Hopper")
    row("uot", "非 NVIDIA 加速器顆數 × 情境係數", "百萬顆", {c_(y): "=(" + "+".join(units_f(g, y) for g in OT) + f")*{umx(y, s)}" for y in BU}, note="AMD、TPU、Trainium、其他 ASIC")
    row("ueff", "有效顆數 ＝ NVIDIA ＋ 權重 × 非 NVIDIA", "百萬顆", {c_(y): f"={c_(y)}{R['unv']}+{P['w']}*{c_(y)}{R['uot']}" for y in BU})
    row("ai0", "FY2026 AI 載板營收（錨點電子事業 × AI 占比）", "¥B", {"D": f"=q26_e*{P['aie']}"}, note="Assumed 切分；公司未揭露")
    row("k", "每顆有效金額（Ibiden 份額 × 載板金額；FY2026 校準）", "千日圓／顆", {"D": f"=D{R['ai0']}/D{R['ueff']}"}, fmt="#,##0.0",
        note="Derived：¥B ÷ 百萬顆＝千日圓／顆。÷ 匯率即美元／顆（Units 頁）")
    row("asp", "每顆載板金額指數（面積、層數；FY2026＝1）", "x", {"D": "=1", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['ga']})" for y in BU[1:]}}, fmt="0.000")
    row("rel", "Ibiden 份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("pp", "短缺溢價指數（FY2026＝1）", "x", {"D": "=1", **{c_(y): f"={P['pp']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("ai", "AI 伺服器載板營收", "¥B", {"D": f"=D{R['ai0']}", **{c_(y): f"={c_(y)}{R['ueff']}*$D${R['k']}*{c_(y)}{R['asp']}*{c_(y)}{R['rel']}*{c_(y)}{R['pp']}" for y in BU[1:]}})
    sec_("B. 電子事業")
    row("ne", "電子事業非 AI 營收（通用伺服器、PC 等）", "¥B", {"D": f"=q26_e-D{R['ai']}", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gne']})" for y in BU[1:]}})
    row("e", "電子事業營收", "¥B", {"C": "=q25_e", **{c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['ne']}" for y in BU}})
    row("mai0", "AI 載板營業利益率（FY2026 校準）", "%", {"D": f"=(q26_eoi-{P['mne']}*D{R['ne']})/D{R['ai']}"}, fmt="0.0%",
        note="Derived：使 FY2026 電子事業營業利益＝錨點；非 AI 營益率為 In 頁 m_ne")
    row("mai", "AI 載板營業利益率 ＝ 1 −（1 − 校準值）÷ 溢價指數", "%", {"D": f"=D{R['mai0']}", **{c_(y): f"=1-(1-$D${R['mai0']})/{c_(y)}{R['pp']}" for y in BU[1:]}}, fmt="0.0%",
        note="Derived：每單位成本不變，價格隨溢價指數升降，差額全部落在利潤")
    row("eoi", "電子事業營業利益", "¥B", {"C": "=q25_eoi", **{c_(y): f"={c_(y)}{R['ai']}*{c_(y)}{R['mai']}+{c_(y)}{R['ne']}*{P['mne']}" for y in BU}})
    row("em", "電子事業營益率", "%", {c_(y): f"={c_(y)}{R['eoi']}/{c_(y)}{R['e']}" for y in FYS}, fmt="0.0%")
    sec_("C. 陶瓷、其他事業與全公司")
    row("cer", "陶瓷事業營收", "¥B", {"C": "=q25_cer", "D": "=q26_cer", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gcer']})" for y in BU[1:]}})
    row("oth", "其他事業營收", "¥B", {"C": "=q25_oth", "D": "=q26_oth", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['goth']})" for y in BU[1:]}})
    row("rev", "營收", "¥B", {c_(y): f"={c_(y)}{R['e']}+{c_(y)}{R['cer']}+{c_(y)}{R['oth']}" for y in FYS})
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI 伺服器載板 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in BU}, fmt="0.0%")
    row("ceroi", "陶瓷＋其他事業營業利益＋調整", "¥B", {"C": "=q25_oi-q25_eoi", "D": "=q26_oi-q26_eoi",
        **{c_(y): f"={c_(y)}{R['cer']}*m_cer+{c_(y)}{R['oth']}*m_oth+q26_adj" for y in BU[1:]}})
    row("oi", "營業利益", "¥B", {c_(y): f"={c_(y)}{R['eoi']}+{c_(y)}{R['ceroi']}" for y in FYS})
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("ni", "淨利（歸屬母公司）", "¥B", {"C": "=q25_ni", "D": "=q26_ni", **{c_(y): f"={c_(y)}{R['oi']}*{P['nioi']}" for y in BU[1:]}})
    row("sh", "股數（分割後）", "十億股", {"D": "=q26_sh", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['shg']})" for y in BU[1:]}}, fmt="0.0000")
    row("eps", "EPS（日圓）", "¥", {"C": "=q25_eps", **{c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh']}" for y in BU}})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — Ibiden 損益（¥B；EPS 日圓）。FY2025A 實績、FY2026 錨點（公司上修後預測），FY2027 起為模型。FY(t)＝t/4–(t+1)/3，對照 CY(t) 顆數。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始）"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A：基準（顆數係數 1.0）", "B": "情境 B：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "說明"])
sr = 4; SUMR = {}
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in (("rev", "營收", "¥B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 伺服器載板", "¥B", "#,##0.0"), ("e", "  電子事業", "¥B", "#,##0.0"),
                                ("em", "電子事業營益率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"), ("ni", "淨利", "¥B", "#,##0.0"),
                                ("eps", "EPS", "¥", "#,##0.0"), ("epsg", "  EPS 年增", "%", "0.0%")):
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS:
            if key in ("g", "epsg") and y == 2025: continue
            if key == "ai" and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 錨點（校準恆等式）").font = BOLD; sr += 1
CHK = {}
for k, lab, ref in (("rev", "營收差（¥B）", "q26_rev"), ("oi", "營業利益差（¥B）", "q26_oi"), ("eps", "EPS 差（日圓）", "q26_eps")):
    PL.cell(sr, 1, f"  情境 A：{lab}"); PL[f"D{sr}"] = f"=D{PB['A'][k]}-{ref}"; PL[f"D{sr}"].number_format = "0.0000"; CHK[k] = sr; sr += 1
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 12; PL.column_dimensions["H"].width = 70
for c in "CDEFG": PL.column_dimensions[c].width = 14

# ---------------- Units：份額與每顆金額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — Ibiden 在 AI 加速器載板的有效金額與份額（本專案權威，規則第 13 條）。顆數取自 markets/network（情境 A）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "FY2026（CY2026）", "FY2027", "FY2028", "FY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("unv", "NVIDIA 加速器顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['unv']}", "#,##0.00", "markets/network"),
        ("uot", "非 NVIDIA 加速器顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['uot']}", "#,##0.00", ""),
        ("ueff", "有效顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['ueff']}", "#,##0.00", "非 NVIDIA × 權重 w_ot"),
        ("kusd", "每顆 NVIDIA 加速器的 Ibiden 營收（美元）", "US$", lambda y: f"=PL!$D${A_['k']}*PL!{FC[y]}{A_['asp']}*PL!{FC[y]}{A_['rel']}*PL!{FC[y]}{A_['pp']}*1000/fx", "#,##0", "＝份額 × 每顆載板金額"),
        ("aspusd", "每顆 NVIDIA 載板金額（美元；假設 FY2026 份額 70%）", "US$", lambda y: f"=PL!$D${A_['k']}*PL!{FC[y]}{A_['asp']}*PL!{FC[y]}{A_['pp']}*1000/fx/0.7", "#,##0", "BofA：Ibiden 占 GPU 載板 70–80%（Interested-party）；先驗比對推導 $150–300"),
        ("ai", "AI 伺服器載板營收", "¥B", lambda y: f"=PL!{FC[y]}{A_['ai']}", "#,##0.0", ""),
        ("aiusd", "AI 伺服器載板營收（美元）", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx", "#,##0.00", "÷ 匯率 152"),
        ("itr", "AI 載板營收 ÷ AI IT 資本支出", "%", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx/itc_{y % 100}", "0.000%", "× 每 GW IT capex 即每 GW 營收"),
        ("pgw", "每 GW Ibiden 營收（VR200 每 GW IT capex 37.59B）", "US$M", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx/itc_{y % 100}*37.59*1000", "#,##0", "Tokenomics IF_CapexIT（VR200 基準，經橋接表）")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：每顆 NVIDIA 加速器 Ibiden 只分到約 1–2 百美元，每 GW 約數千萬美元；瓶頸反映在利潤率（電子事業營益率）與倍數，不在營收規模。").font = BOLD
UNOTE = ur
UN.column_dimensions["A"].width = 60; UN.column_dimensions["B"].width = 10
for c in "CDEF": UN.column_dimensions[c].width = 16
UN.column_dimensions["G"].width = 70

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（日圓）。結論用「現價 ÷ FY2026–FY2028 EPS（模型）」與 what must be true。十年 P/E 規則值待 Bloomberg（10-27 抓取單加 4062）：公開資料商的 Ibiden 十年 P/E 本版未取得，規則值本版不計"
V["A2"] = "循環位置：Ibiden 不套設備股見頂年；以本模型營收最高的會計年度為見頂年（比照 NVIDIA、Arista）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
V.cell(4, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, 5, ["項目", "A", "B", "C", "說明"])
V.cell(6, 1, "現價（2026-10-23 收盤，分割後）"); V["B6"] = "=px"; V["B6"].number_format = "#,##0"
V.cell(7, 1, "現價 ÷ FY2026 EPS（錨點）"); V.cell(8, 1, "現價 ÷ FY2027 EPS（模型）"); V.cell(9, 1, "現價 ÷ FY2028 EPS（模型）"); V.cell(10, 1, "現價 ÷ FY2029 EPS（模型）")
for j, s in enumerate(SC):
    c = "BCD"[j]
    for rr, y in ((7, 2026), (8, 2027), (9, 2028), (10, 2029)):
        V[f"{c}{rr}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{rr}"].number_format = "0.0"
V.cell(11, 1, "前瞻 P/E（stockanalysis）"); V["B11"] = "=pe_fwd_sa"; V["B11"].number_format = "0.0"
V.cell(12, 1, "  隱含未來 12 月 EPS ＝ 取得時價格（10-16 9,764）÷ 前瞻 P/E"); V["B12"] = "=cons_feps"; V["B12"].number_format = "0.0"
V.cell(13, 1, "近 12 月 P/E（stockanalysis）"); V["B13"] = "=pe_ttm"; V["B13"].number_format = "0.0"
V["E8"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"; V["E11"] = "口徑：未來 12 個月"
V.cell(15, 1, "2. 循環位置（本模型營收最高年）").font = BOLD
hdr(V, 16, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "見頂年是否未知", "說明"])
for i, s in enumerate(SC, 17):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s)
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 5, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年；對應倍數待十年 P/E" if s == "A" else "同上")
V.cell(21, 1, "3. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（看 FY2028 EPS）").font = BOLD
hdr(V, 22, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：Andy 合理 P/E"])
vr = 23
for x in (0.5, 0.75, 1.0, 1.21, 1.5):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_fair_andy").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數愈高，現價愈依賴「短缺溢價延續到 2028 年以後」。"); WMT_NOTE = vr
vr += 2
V.cell(vr, 1, "4. 對照：以 Andy 合理 P/E 25x 計算的隱含價（只作對照，非規則值）").font = BOLD; vr += 1
hdr(V, vr, ["情境", "FY2027 EPS", "× 25x", "相對現價", "FY2028 EPS", "× 25x", "相對現價"]); IMPH = vr; vr += 1
IMP = {}
for s in SC:
    V.cell(vr, 1, s)
    V.cell(vr, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.0"; V.cell(vr, 3, f"=B{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 4, f"=C{vr}/px-1").number_format = "0%"
    V.cell(vr, 5, f"=PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"; V.cell(vr, 6, f"=E{vr}*pe_fair_andy").number_format = "#,##0"; V.cell(vr, 7, f"=F{vr}/px-1").number_format = "0%"
    for c in range(1, 8): V.cell(vr, c).fill = KEY
    IMP[s] = vr; vr += 1
V.column_dimensions["A"].width = 58
for c in "BCDEFG": V.column_dimensions[c].width = 18

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；每顆有效金額與 AI 載板營益率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 載板營收、電子事業營益率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 載板", "FY2027 電子營益率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 載板", "FY2028 電子營益率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("溢價快速回落", "價格溢價 2027 0.95、2028 0.85、2029 0.75（新產能提早到位）", dict(pp="PPLO")),
         ("溢價延續", "價格溢價 2027 1.05、2028 1.05、2029 1.00", dict(pp="PPHI")),
         ("份額流失（低）", "份額相對值 2027 0.85、2028 0.75、2029 0.70（欣興、三星電機加速）", dict(rel="RLO")),
         ("份額持平（高）", "份額相對值 1.00", dict(rel="RHI")),
         ("每顆金額年增 10%", "面積、層數成長較慢", dict(ga="g_asp_lo")), ("每顆金額年增 22%", "較快", dict(ga="g_asp_hi")),
         ("非 NVIDIA 權重 0.2", "ASIC 多由其他載板廠供應", dict(w="w_ot_lo")), ("非 NVIDIA 權重 0.6", "Ibiden 在 ASIC 份額較高", dict(w="w_ot_hi")),
         ("AI 占電子事業 55%", "FY2026 切分較低（校準跟著變）", dict(aie="ai_e_lo")), ("AI 占電子事業 75%", "較高", dict(aie="ai_e_hi")),
         ("非 AI 營益率 10%", "校準後 AI 載板營益率較高", dict(mne="m_ne_lo")), ("非 AI 營益率 20%", "較低", dict(mne="m_ne_hi")),
         ("淨利 ÷ 營業利益 0.62", "稅與非營業較不利", dict(nioi="ni_oi_lo")),
         ("CB 轉換：股數 +3%／年", "2031 年到期 CB 轉換稀釋", dict(shg="sh_g_hi"))]
for nm, base, suf in (("PPLO", "pp", "_lo"), ("PPHI", "pp", "_hi"), ("RLO", "rel", "_lo"), ("RHI", "rel", "_hi")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}_{y}"] = DefinedName(f"{nm}_{y}", attr_text=wb.defined_names[f"{base}_{y}{suf}"].attr_text)
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "ai", 2027, "#,##0.0"), ("E", "em", 2027, "0.0%"), ("F", "eps", 2027, "#,##0.0"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "ai", 2028, "#,##0.0"), ("J", "em", 2028, "0.0%"), ("K", "eps", 2028, "#,##0.0")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
assert nn + 3 < 40
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 三個顆數情境相同；B、C 的差異從 FY2028 起。溢價指數同時改變營收與 AI 載板營益率（成本不變），所以對 EPS 的影響大於對營收的影響。"); SENS_NOTE = nn + 3
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司目標、研調份額、先驗。只作事後對照，不為對上數字而調參數。FY2027 以後的共識：未找到（付費）"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.0", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 公司目標與市場（校準未使用）").font = BOLD; rr_ += 1
rec("FY2029 營收 vs 公司 FY2030 目標 ¥1 兆", "=tgt30_rev", pl("rev", 2029), "Quartr 摘要（Interested-party）", "模型到 FY2029；目標年晚一年")
rec("FY2029 營業利益 vs 公司 FY2030 目標 ¥3,000 億", "=tgt30_oi", pl("oi", 2029), "同上", "")
rec("FY2027 EPS vs 前瞻 P/E 隱含未來 12 月 EPS", "=cons_feps", pl("eps", 2027), "stockanalysis（推導）", "未來 12 月 ≈ ½ FY2026＋½ FY2027，口徑不同，只作量級比較")
RC.cell(rr_, 1, "B. 份額與每顆金額（先驗比對的機制檢查）").font = BOLD; rr_ += 1
rec("FY2026 每顆 NVIDIA 載板金額（美元）vs 先驗比對推導中點 $225", 225, lambda s: f"=Units!C{UR['aspusd']}", "先驗比對_Ibiden（Drive）", "推導區間 $150–300；封裝全部（中介層＋載板＋組裝）B200 約 $1,100（Silicon Analysts）", fmt="#,##0")
rec("FY2026 每 GW Ibiden 營收（US$M）vs 先驗比對推導中點 75", 75, lambda s: f"=Units!C{UR['pgw']}", "先驗比對_Ibiden（Drive）", "", fmt="#,##0")
rec("FY2027 AI 載板 ÷ 營收", None, pl("aish", 2027), "—", "", fmt="0.0%")
rec("FY2028 電子事業營益率 vs FY2026 錨點（百分點）", "=q26_eoi/q26_e", pl("em", 2028), "錨點", "短缺溢價回落使營益率下滑", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "C. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "產能上限（Ibiden 2028 年為 2024 年 2.5–2.8 倍）與需求面積的對照：2024 年基準面積未找到，本版不計"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以 AI 載板營收 ÷ 部署 GW 取代 Units 頁的每 GW 換算"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 70

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝決算短信與補足資料；分部門預測取自 Q1 補足資料的轉載）"; SO["A1"].font = BOLD
srcs = [("公司：2026 年 3 月期決算短信（FY2025 實績、分部門）", TAN25_URL),
        ("公司：2027 年 3 月期 Q1 決算短信（上修後全年預測）", TANQ1_URL),
        ("公司：Q1 補足資料（分部門全年預測、設備投資 2,100 億、匯率 ¥152）", SUPQ1_URL),
        ("二手：SEMICON.TODAY Q1 深掘（分部門 Q1、EPS 分割後 149.68）", "https://semicon.today/archives/12352"),
        ("二手：LIMO Q1（預收款、股價）", "https://finance.limo.media/articles/detail/683"),
        ("二手：Quartr FY2025 摘要（FY2030 目標；AI 摘要，營收數字有誤處未採用）", "https://quartr.com/events/ibiden-co-ltd-4062-q4-2026_F9Xvd7IW"),
        ("研調：Tom's Hardware ABF 2026（供需缺口、Ibiden 路線圖、競爭者擴產、ABF 膜漲價）", "https://www.tomshardware.com/tech-industry/semiconductors/the-state-of-abf-substrates-in-data-center-silicon-in-2026-solving-the-supply-crunch-and-material-wall-beneath-every-ai-accelerator"),
        ("研調：Data Gravity（高階 FC-BGA 份額、ABF 載板平均單價）", "https://www.datagravity.dev/p/the-abf-substrate-bottleneck"),
        ("研調：thebuildout（BofA：Ibiden 占 GPU 載板 70–80%、短缺到 2028）", "https://thebuildout.ai/intel/229c44f2-da57-4f46-b5fa-e5d03702236c"),
        ("研調：Encyclopedia of Stocks（欣興 2025 年 NVIDIA 約 30%）", "https://encyclopediaofstocks.substack.com/p/abf-substrate-market-ai-driven-super"),
        ("研調：Silicon Analysts（封裝成本，含載板）", "https://siliconanalysts.com/data/ai-chip-costs"),
        ("市場：stockanalysis（10-16 股價、前瞻與近 12 月 P/E、共識目標價；同業前瞻 P/E、ROE）", "https://stockanalysis.com/quote/tyo/4062/"),
        ("市場：10-23 收盤（markets/invest_map/data/prices.csv：株探、松井證券）", "https://kabutan.jp/stock/?code=4062"),
        ("市場層：markets/network、tsmc 成品（本 repo，唯讀）", ""),
        ("待補（人工）：十年 P/E 歷史（Bloomberg）；AI 伺服器載板營收（公司未揭露）；FY2027 以後共識（付費）；CB 發行額；ASIC 載板份額", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 90; SO.column_dimensions["B"].width = 140

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("Ibiden（4062）收支與估值模型 v0.1.1（AI 半導體 Project：載板與 PCB 樣板）　2026-10-25", True),
 ("版本：v0.1（初版）。Ibiden 在 AI 加速器載板的有效金額與份額在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）。", False),
 ("模型命題：Ibiden 的 AI 伺服器載板營收＝AI 加速器顆數（本專案網通模型）× Ibiden 份額 × 每顆載板金額（面積、層數逐年增加）× 短缺溢價；電子事業的非 AI 部分、陶瓷與其他事業另依成長率推。股價約 54–65 倍前瞻 EPS：2028 年產能擴到 2.5 倍、競爭者同時擴產之後，短缺溢價能否守住電子事業約 29% 的營益率？", False),
 ("核心問題：短缺溢價回落多少、何時回落，FY2028 EPS 會少多少；現價隱含的溢價路徑合不合理。", False), ("", False),
 ("會計年度 4 月至隔年 3 月：FY(t)＝t/4–(t+1)/3，對照 CY(t) 的加速器顆數（重疊 9 個月，不平移）。預測期 FY2027–FY2029（三年）。金額：十億日圓；EPS：日圓（2026-10-01 一拆二後）。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("AI 載板：markets/network 的加速器顆數（NVIDIA ＋ 權重 × 非 NVIDIA）× 每顆有效金額（FY2026 校準）× 金額指數（+18%／年，Analogy）× 份額相對值（Assumed）× 溢價指數（Assumed）。", False),
 ("利潤率：AI 載板營益率 FY2026 校準（非 AI 部分 15%，Assumed），之後＝ 1 −（1 − 校準值）÷ 溢價指數（成本不變）。陶瓷、其他事業營益率沿用 FY2026 預測。淨利 ÷ 營業利益 0.661（FY2026 預測）。", False),
 ("校準：FY2026 ＝ 公司 2026-08-04 上修後預測（營收 5,500 億、營業利益 1,270 億、淨利 840 億、EPS 149.68）；AI 伺服器載板占電子事業 65%（Assumed，公司未揭露）。", False),
 ("估值：現價 ÷ FY2026–FY2028 EPS 與 what must be true；十年 P/E 規則值待 Bloomberg（本版不計）；循環位置以本模型營收最高年判定。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。", False),
 ("", False),
 ("已知限制：(1) AI 載板營收公司未揭露，以 Assumed 切分校準；(2) 溢價指數與份額相對值皆為 Assumed，是 FY2028 EPS 最敏感的輸入（Sens）；(3) 未設產能上限（2024 基準面積未找到）；(4) FY2026 錨點為公司預測，Q2 決算後更新；(5) 十年 P/E 規則值待 Bloomberg；(6) CB 轉換稀釋只在 Sens；(7) R4 產業部署路徑未建（Recon 預留列）。", False)]
lines.append(("v0.1.1（2026-10-25）：Val 頁新增「0. 目標價與上檔空間」與目標價用的十年 P/E（companiesmarketcap，Interested-party，單一來源）（放在 Val 既有內容之後，不位移既有儲存格），In 頁新增目標價情境權重 tp_wA／tp_wB／tp_wC（0.3／0.5／0.2，Assumed）。目標價＝十年 P/E 中位數 × 各情境 FY2027 EPS；其餘儲存格不變。", False))
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ---------------- Val 第 0 節：目標價與上檔空間（v0.1.1 新增；Andy 2026-10-25 要求，列標籤依工作單逐字） ----------------
# 為不位移既有儲存格（與舊版 dist 逐格相同），本節放在 Val 既有內容之後；HTML 估值章第一個區塊即為本節
def add_target_price(V, imp, note, fmt, method, px_note, t0=None):
    t0 = t0 or V.max_row + 2
    V.cell(t0, 1, "舊法：規則值情境加權（參考；v0.1.1 的「0. 目標價與上檔空間」改名，不再是目標價）").font = BOLD
    hdr(V, t0 + 1, ["項目", "數值", "", "", "說明"])
    V.cell(t0 + 2, 1, "舊法隱含價（情境加權，參考）"); V[f"B{t0 + 2}"] = f"=tp_wA*B{t0 + 4}+tp_wB*B{t0 + 5}+tp_wC*B{t0 + 6}"; V[f"B{t0 + 2}"].number_format = fmt
    V[f"E{t0 + 2}"] = "權重 × 各情境隱含價"
    V.cell(t0 + 3, 1, "舊法隱含價相對現價（÷ 現價 − 1）"); V[f"B{t0 + 3}"] = f"=B{t0 + 2}/px-1"; V[f"B{t0 + 3}"].number_format = "0.0%"; V[f"E{t0 + 3}"] = px_note
    for i, (s, lab) in enumerate((("A", "情境 A 隱含價（樂觀）"), ("B", "情境 B 隱含價（基準）"), ("C", "情境 C 隱含價（下檔）"))):
        rr = t0 + 4 + i
        V.cell(rr, 1, lab); V[f"B{rr}"] = imp[s]; V[f"B{rr}"].number_format = fmt; V[f"E{rr}"] = note[s]
    rr = t0 + 7
    V.cell(rr, 1, "情境權重 A／B／C"); V[f"B{rr}"] = "=tp_wA"; V[f"C{rr}"] = "=tp_wB"; V[f"D{rr}"] = "=tp_wC"
    for c in "BCD": V[f"{c}{rr}"].number_format = "0%"
    V[f"E{rr}"] = f'="合計 "&TEXT(B{rr}+C{rr}+D{rr},"0%")&"；In 頁 tp_wA／tp_wB／tp_wC（Assumed，Andy 可改）"'
    V.cell(t0 + 8, 1, "舊法說明"); V[f"B{t0 + 8}"] = method
    return t0
TP0 = V.max_row + 2
# 目標價用十年 P/E（本節下方）：companiesmarketcap 年底 P/E（2026-10-25 取得）；2017 為虧損（負值），不計入中位數
PE_CMC = [(2016, 44.1), (2017, -6.66), (2018, 35.3), (2019, 137.0), (2020, 55.2), (2021, 53.1), (2022, 40.7), (2023, 26.8), (2024, 57.2), (2025, 33.6)]
pr = TP0 + 10
V.cell(pr, 1, "0a. 十年 P/E（舊法用；companiesmarketcap 日曆年年底 P/E，Interested-party，資料商；第二來源：未找到）").font = BOLD
hdr(V, pr + 1, ["年底", "P/E", "計入中位數", "說明"])
for i, (y, v) in enumerate(PE_CMC):
    rw = pr + 2 + i
    V.cell(rw, 1, str(y)); V.cell(rw, 2, v).font = BLUE; V.cell(rw, 2).number_format = "0.0"
    V.cell(rw, 3, "否" if v <= 0 else "是")
    V.cell(rw, 4, "https://companiesmarketcap.com/ibiden/pe-ratio/（2026-10-25 讀取）" if i == 0 else ("虧損年，P/E 為負，不計" if v <= 0 else ""))
PE_MED_ROW = pr + 12
V.cell(PE_MED_ROW, 1, "中位數（計入者）"); V[f"B{PE_MED_ROW}"] = "=MEDIAN(" + ",".join(f"B{pr + 2 + i}" for i, (y, v) in enumerate(PE_CMC) if v > 0) + ")"
V[f"B{PE_MED_ROW}"].number_format = "0.0"
add_target_price(V, {s: f"=$B${PE_MED_ROW}*PL!E{SUMR[(s, 'eps')]}" for s in SC},
    {s: f"情境 {s} FY2027 EPS（PL 頁）× 十年 P/E 中位數（第 0a 節）" for s in SC}, "#,##0",
    "P/E：十年（2016–2025 日曆年年底）P/E 中位數（companiesmarketcap 單一來源，第 0a 節；2017 年虧損不計，9 年）× 各情境 FY2027 EPS（FY2027＝2027-04～2028-03）。"
    "本模型無規則值（十年 P/E 待 Bloomberg），依工作單改用此備案；倍數為日曆年年底口徑，與 3 月制會計年度不完全對齊。十年倍數多為 AI 載板需求前的時期（2019 年獲利低谷約 137x 為極值，中位數受影響小），偏誤方向不明確；"
    "FY2027 三情境 EPS 相同（情境差從 FY2028 起），故三情境隱含價相同；隱含價約為 FY2027 年底（2028-03）的價格、未折現，當作 12 個月目標價時偏高。",
    "現價＝In 頁 px（2026-10-23 收盤，分割後）", t0=TP0)

V2["old_row"] = TP0
# ================= v0.2（2026-10-25）：PL 延伸到 FY2032、Val 第 0 節目標價（DCF 與 Comparable 加權）、Sens 目標價參數 =================
# 全部放在各頁既有內容之後，不位移既有儲存格（FY2025–FY2029 數值與 v0.1.1 相同）
FADEFILL = PatternFill("solid", fgColor="E7E6E6")
XY = (2025, 2026, 2027, 2028, 2029, 2030, 2031, 2032); XC = {y: "CDEFGHIJ"[i] for i, y in enumerate(XY)}; FADE = (2030, 2031, 2032)
K2 = V2["keys"]
for rw in range(4, IN.max_row + 1):
    if IN.cell(rw, 1).value in ("tp_wA", "tp_wB", "tp_wC", "tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "nim10_f"):
        IN.cell(rw, 5).fill = KEY
# --- PL 延伸表（三情境到 FY2032） ---
xr = PL.max_row + 3; XT = xr
PL.cell(xr, 1, "延伸：三情境到 FY2032（v0.2 新增）。FY2025–FY2029 引用上方推導區塊；FY2030 起為收斂法（灰底），非驅動因子推導：營收年增由 FY2029 線性收斂到 fade_g、淨利率由 FY2029 線性收斂到收斂目標，FY2032 達到").font = BOLD
xr += 1; XR = {}; XH = None
for s in SC:
    E = PB[s]; R = {}
    PL.cell(xr, 1, f"延伸｜{SCN_T[s]}").font = BOLD; xr += 1
    hdr(PL, xr, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
    XH = XH or xr; xr += 1
    for key, lab, unit in (("rev", "營收（延伸）", V2["unit"]), ("g", "  營收年增", "%"), ("ni", "淨利（延伸）", V2["unit"]), ("nim", "  淨利率", "%"),
                           ("sh", "股數（延伸）", V2["shunit"]), ("eps", "EPS（延伸）", V2["cur"]), ("epsg", "  EPS 年增", "%"),
                           ("tg", "收斂目標：FY2032 營收年增", "%"), ("tm", "收斂目標：FY2032 淨利率", "%")):
        PL.cell(xr, 1, lab); PL.cell(xr, 2, unit); R[key] = xr; xr += 1
    G = lambda k: f"$G${R[k]}"
    for y in XY:
        c = XC[y]; pc = XC.get(y - 1)
        if y <= 2029:
            PL[f"{c}{R['rev']}"] = f"={c}{E[K2['rev']]}"
            if y > 2025: PL[f"{c}{R['g']}"] = f"={c}{R['rev']}/{pc}{R['rev']}-1"
            PL[f"{c}{R['ni']}"] = f"={c}{E[K2['ni']]}"
            if y > 2025 or K2.get("sh25", True): PL[f"{c}{R['sh']}"] = f"={c}{E[K2['sh']]}"
            PL[f"{c}{R['eps']}"] = f"={c}{E[K2['eps']]}"
        else:
            k = y - 2029
            PL[f"{c}{R['g']}"] = f"={G('g')}+($C${R['tg']}-{G('g')})*{k}/3"
            PL[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
            PL[f"{c}{R['nim']}"] = f"={G('nim')}+($C${R['tm']}-{G('nim')})*{k}/3"
            PL[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
            PL[f"{c}{R['sh']}"] = f"={pc}{R['sh']}*(1+{V2['shg']})"
            PL[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh']}"
        if y <= 2029: PL[f"{c}{R['nim']}"] = f"={c}{R['ni']}/{c}{R['rev']}"
        if y > 2025: PL[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("rev", "#,##0.0"), ("g", "0.0%"), ("ni", "#,##0.0"), ("nim", "0.0%"), ("sh", "0.0000"), ("eps", V2["epsfmt"]), ("epsg", "0.0%")):
            PL[f"{c}{R[key]}"].number_format = fm
            if y in FADE: PL[f"{c}{R[key]}"].fill = FADEFILL
    PL[f"C{R['tg']}"] = f"=fade_g{s}"; PL[f"C{R['tg']}"].number_format = "0.0%"
    PL[f"C{R['tm']}"] = {"A": f"=G{R['nim']}", "B": f"=G{R['nim']}*(1+nim10_f)/2", "C": f"=G{R['nim']}*nim10_f"}[s]; PL[f"C{R['tm']}"].number_format = "0.0%"
    PL.cell(R["rev"], 11, "FY2030 起為收斂法，非驅動因子推導（上游市場層只到 2029）")
    PL.cell(R["tg"], 11, {"A": "In 頁 fade_gA＝tp_g＋2pt", "B": "In 頁 fade_gB＝tp_g", "C": "In 頁 fade_gC＝tp_g−2pt"}[s])
    PL.cell(R["tm"], 11, {"A": "＝本情境 FY2029 淨利率", "B": "＝FY2029 淨利率與十年平均替代值（FY2029 × nim10_f）的中點", "C": "＝十年平均替代值（FY2029 × nim10_f；十年平均未找到）"}[s])
    PL.cell(R["sh"], 11, f"FY2030 起股數年變動＝{V2['shg']}")
    XR[s] = R; xr += 1
PL.column_dimensions["I"].width = 14; PL.column_dimensions["J"].width = 14; PL.column_dimensions["K"].width = 60
# --- Val 第 0 節：目標價 ---
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
t0 = V.max_row + 3; T0 = t0
V.cell(t0, 1, "0. 目標價（DCF 與 Comparable 加權；v0.2 新增，放在本頁既有內容之後以免位移既有儲存格；三情境欄 B／C／D＝A 順利／B 平平／C 不如預期）").font = BOLD
V.cell(t0 + 1, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, t0 + 2, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"]); T["h1"] = t0 + 2
r_ = t0 + 3
V.cell(r_, 1, "情境 EPS（FY2027–FY2032）"); V[f"E{r_}"] = "PL 頁延伸表；FY2030 起為收斂法"; T["eps_h"] = r_; r_ += 1
EPSR = {}
for y in YRS6:
    tset(r_, f"eps{y}", f"  FY{y}", [f"=PL!{XC[y]}{XR[s]['eps']}" for s in SC], V2["epsfmt"]); EPSR[y] = r_; r_ += 1
V.cell(r_, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = r_; r_ += 1
tset(r_, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）"); RV = r_; r_ += 1
CFR = {}
for y in YRS6:
    tset(r_, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B${RV})" for c in "BCD"], V2["epsfmt"]); CFR[y] = r_; r_ += 1
def dcf_formula(c, r="tp_r", g="tp_g", reinv=None):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    reinv = reinv or f"Val!$B${RV}"
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
c1, c6 = CFR[2027], CFR[2032]
tset(r_, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}{c1}:{c}{c6},1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}{c6}*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], V2["pxfmt"],
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底，t＝1 對 FY2027"); r_ += 1
tset(r_, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}{T['dcfv']}*(1+tp_r)" for c in "BCD"], V2["pxfmt"]); r_ += 1
tset(r_, "dcftp", "DCF 目標價（情境加權）", [f"=tp_wA*B{T['dcf12']}+tp_wB*C{T['dcf12']}+tp_wC*D{T['dcf12']}", None, None], V2["pxfmt"], "Σ tp_wA／wB／wC × 情境價值 ×(1＋tp_r)；權重見 In 頁"); r_ += 2
V.cell(r_, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD; r_ += 1
hdr(V, r_, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"]); T["ph"] = r_; r_ += 1
P1 = r_
for nm, tk, pe, dt in V2["peers"]:
    V.cell(r_, 1, nm); V.cell(r_, 2, tk); cc = V.cell(r_, 3, pe); cc.font = BLUE; cc.number_format = "0.0"
    V.cell(r_, 4, dt); V.cell(r_, 5, "stockanalysis.com 個股頁 Forward PE（分析師共識；價格日期如左）"); V.cell(r_, 6, "Interested-party（賣方共識）"); r_ += 1
P2 = r_ - 1; T["p1"], T["p2"] = P1, P2
tset(r_, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C{P1}:C{P2})", None, None], "0.0", V2["peer_note"]); r_ += 1
w1, w2 = V2["fwd_w"]
tset(r_, "feps", "12 個月後的前瞻 EPS（情境加權）", ["=" + "+".join(f"tp_w{s}*({w1}*{c}{EPSR[2027]}+{w2}*{c}{EPSR[2028]})" for s, c in zip(SC, "BCD")), None, None], V2["epsfmt"], V2["fwd_note"]); r_ += 1
tset(r_, "comptp", "Comparable 目標價", [f"=B{T['pmed']}*B{T['feps']}", None, None], V2["pxfmt"], "同業中位數 × 上列 EPS"); r_ += 2
V.cell(r_, 1, "0.3 結論").font = BOLD; r_ += 1
hdr(V, r_, ["項目", "數值", "", "", "說明"]); T["h3"] = r_; r_ += 1
tset(r_, "tpw", "目標價（DCF 與 Comparable 加權）", [f"=tp_w_dcf*B{T['dcftp']}+(1-tp_w_dcf)*B{T['comptp']}", None, None], V2["pxfmt"], "tp_w_dcf × DCF 目標價＋（1 − tp_w_dcf）× Comparable 目標價；tp_w_dcf 見 In 頁"); r_ += 1
tset(r_, "comp2", "Comparable 目標價（市場主流同業法）", [f"=B{T['comptp']}", None, None], V2["pxfmt"], "引用 0.2"); r_ += 1
tset(r_, "cons", "市場共識目標價", ["=cons_tp", None, None], V2["pxfmt"], V2["cons_tp"][1]); r_ += 1
tset(r_, "upw", "上檔空間：加權目標價", [f"=B{T['tpw']}/px-1", None, None], "0.0%", V2["px_note"]); r_ += 1
tset(r_, "upc", "上檔空間：Comparable 目標價", [f"=B{T['comp2']}/px-1", None, None], "0.0%"); r_ += 1
tset(r_, "geps", "與市場共識目標價的差距：EPS 因素", [f"=B{T['feps']}/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS 見 In 頁 cons_feps 與 Recon"); r_ += 1
tset(r_, "gmul", "與市場共識目標價的差距：倍數因素", [f"=(cons_tp/cons_feps)/B{T['pmed']}-1", None, None], "0.0%"); r_ += 1
V[f"E{T['gmul']}"] = f'="共識目標價 ÷ 共識前瞻 EPS 隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B{T["pmed"]},"0.0")&"x"'
tset(r_, "band", "價格帶（DCF 不如預期～順利）", [f'=TEXT(D{T["dcf12"]},"#,##0")&"～"&TEXT(B{T["dcf12"]},"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）"); r_ += 1
PEROW = r_ + 3
PE10 = f"$B${PEROW + 1}:$K${PEROW + 1}"
tset(r_, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）",
     [f"=IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E 分布（下方列）的分位；低於最小值記 0%、高於最大值記 100%；只作檢查，不進目標價"); r_ += 1
V.cell(r_, 1, "目標價方法"); V[f"B{r_}"] = V2["method"]; T["m"] = r_; r_ += 1
V.cell(PEROW, 1, "十年年底 P/E（歷史分位檢查用；" + V2["pe10"][1] + "）").font = BOLD
for j, (yy, vv) in enumerate(V2["pe10"][0]):
    V.cell(PEROW, 2 + j, str(yy))
    cc = V.cell(PEROW + 1, 2 + j, vv); cc.font = BLUE; cc.number_format = "0.0"
V.cell(PEROW + 1, 1, "年底 P/E"); T["pe10h"] = PEROW; T["pe10"] = PEROW + 1
for k_ in ("dcftp", "comptp", "tpw"):
    V[f"B{T[k_]}"].fill = KEY
V.cell(T["tpw"], 1).font = BOLD
defname("tp_final", f"Val!$B${T['tpw']}"); defname("tp_dcf", f"Val!$B${T['dcftp']}"); defname("tp_comp", f"Val!$B${T['comptp']}")
V["A3"] = f"目標價：見第 0 節（本頁第 {T0} 列；v0.2 新增，DCF 與 Comparable 加權）。第 {V2['old_row']} 列「舊法：規則值情境加權（參考）」不再是目標價，只作參考"
for c in "IJK": V.column_dimensions[c].width = 10
# --- Sens：目標價參數 ---
TPS0 = SE.max_row + 3
SE.cell(TPS0, 1, "目標價參數（v0.2）：一次動一個，對加權目標價的影響（只改 DCF 參數；PL 收斂目標 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = V2["pxfmt"]
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "tp_r 10%、tp_g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "tp_r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "tp_r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "tp_g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "tp_g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.4", "tp_w_dcf 0.4", dict(w="0.4")), ("DCF 權重 0.8", "tp_w_dcf 0.8", dict(w="0.8"))]
for k_, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k_, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
# --- Recon：目標價對帳 ---
rq = RC.max_row + 2
RC.cell(rq, 1, "E. 目標價對帳（v0.2；模型欄為情境加權值，非情境 A）").font = BOLD; rq += 1
RECV2 = rq
for lab, ref, mdl, src in (("12 個月後的前瞻 EPS（情境加權）vs 共識前瞻 EPS", "=cons_feps", f"=Val!B{T['feps']}", V2["cons_feps"][1]),
                           ("目標價（DCF 與 Comparable 加權）vs 市場共識目標價", "=cons_tp", f"=Val!B{T['tpw']}", V2["cons_tp"][1]),
                           ("Comparable 目標價 vs 市場共識目標價", "=cons_tp", f"=Val!B{T['comptp']}", "同上")):
    RC.cell(rq, 1, lab); RC.cell(rq, 2, ref).number_format = V2["epsfmt"]; RC.cell(rq, 3, mdl).number_format = V2["epsfmt"]
    RC.cell(rq, 4, f"=C{rq}/B{rq}-1").number_format = "0.0%"; RC.cell(rq, 9, src); RC.cell(rq, 10, "共識只作對帳，不作校準"); rq += 1
# --- README ---
for rw in range(1, RD.max_row + 1):
    v_ = RD.cell(rw, 1).value
    if isinstance(v_, str) and rw == 1: RD.cell(rw, 1, v_.replace("v0.1.1", "v0.2"))
    elif isinstance(v_, str) and v_.startswith("估值："): RD.cell(rw, 1, V2["readme_val"] + "（v0.1.1 原估值說明，保留作參考：" + v_[3:] + "）")
    elif isinstance(v_, str) and v_.startswith("版本："): RD.cell(rw, 1, "版本：v0.2（2026-10-25：目標價改為 DCF 與 Comparable 加權、三情境延伸到 FY2032；FY2025–FY2029 數值不變）。前版：" + v_[3:])
rn = RD.max_row + 1
for t_ in V2["readme_add"]:
    RD.cell(rn, 1, t_); rn += 1
if os.environ.get("V2_ROWMAP"):
    import json as _json
    _json.dump({"XT": XT, "XH": XH, "XR": XR, "T": T, "EPSR": EPSR, "CFR": CFR, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "RECV2": RECV2},
               open(os.environ["V2_ROWMAP"], "w"), ensure_ascii=False, indent=1)

# ---------------- v0.2.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00／14:17）----------------
# 新列一律放在 Val、README 工作表既有內容之後，不位移既有列；既有數值不變
VER_PREV, VER_NOW = 'v0.2', 'v0.2.1'
OBS = [
 "目標價 ¥5,236（DCF 1,931、Comparable 8,541 各半），比現價 8,512（2026-10-23）低 38.5%，級距 −25%～−50%；Comparable 單獨為 +0.3%。主因是 DCF 不是 EPS：本模型 12 個月後前瞻 EPS 202.0 比共識前瞻 EPS 180.9 高 11.7%，但 ROE 近三年平均 8.6% 低於折現率 10%，再投資率 46% 使成長不創造價值，DCF 比現價低 77.3%；Comparable 端同業前瞻 P/E 中位數 42.3 倍、現價 ÷ 共識前瞻 EPS 47.0 倍、共識目標價 12,432 隱含 68.7 倍。ROE 用的是 AI 載板放量前的年度，DCF 對 ABF 漲價路徑的反映有限，待 Andy 判斷。",
 "短缺溢價是最敏感的變數：溢價快速回落（2028 年指數 0.85）時 FY2028 EPS −24.6%、電子事業營益率由 29.1% 降到 23.0%，溢價延續則 +24.6%；份額流失（2028 年相對值 0.75）−11.7%；顆數情境 C −31.6%。",
 "自由現金流（v0.2.1 新增；淨利＋折舊 − 資本支出，不含營運資金變動；折舊、資本支出各取營收的 15%、25%）情境 B 的 FY2027 為 ¥50.4B、FY2028 為 ¥36.2B；若 FY2026 照公司計畫（設備投資 2,100 億、折舊 810 億）則約 −¥45B，固定比率會攤平擴產前段的負值。下一個檢驗點是 FY2026 上半年決算（預計 2026-10 底或 11 月初）：全年預測與設備投資 2,100 億是否調整。"
]   # 重點觀察（數字取自重算後的 Val／PL／Sens；與 README.md 相同文字）
UPW_ROW = next(r_ for r_ in range(1, V.max_row + 1) if V.cell(r_, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 1; REC_ROW = GRADE_ROW + 1
_u = f"B{UPW_ROW}"
V.cell(GRADE_ROW, 1, "上檔空間級距")
V[f"B{GRADE_ROW}"] = f'=IF({_u}>0.5,"大於 50%",IF({_u}>0.25,"25%–50%",IF({_u}>=0,"0%–25%",IF({_u}>=-0.25,"0%～−25%",IF({_u}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"B{GRADE_ROW}"].number_format = "@"
V[f"E{GRADE_ROW}"] = f"以「上檔空間：加權目標價」（第 {UPW_ROW} 列）分六級（Andy 2026-10-25 14:00）"
V.cell(REC_ROW, 1, "投資建議")
V[f"B{REC_ROW}"] = f'=IF({_u}>0.5,"強力買進",IF({_u}>0.25,"買進",IF({_u}>=-0.25,"中立",IF({_u}>=-0.5,"放空","強力放空"))))'
V[f"B{REC_ROW}"].number_format = "@"
V[f"E{REC_ROW}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）"
# --- 自由現金流（算力收支 2026-10-25 請求；地圖依列標籤讀取）：In 頁最後新增折舊 ÷ 營收、資本支出 ÷ 營收；PL 頁最後新增一列（三情境 × 模型年度） ---
r_in = IN.max_row + 2
IN.cell(r_in, 1, "K. 自由現金流假設（v0.2.1 新增；算力收支 2026-10-25 請求；只用於 PL 頁最後一列「自由現金流」，不進 EPS、DCF 與目標價）").font = BOLD; r_in += 1
for key_, item_, lo_, base_, hi_, s1_, s2_ in (
        ("da_rev", "折舊 ÷ 營收", 0.13, 0.15, 0.17,
         "FY2025 實績（2026 年 3 月期）：折舊攤銷 662 億 ÷ 營收 4,162 億＝15.9%（stockanalysis 現金流量表，Verified）",
         "FY2026 公司計畫：折舊 810 億 ÷ 營收預測 5,500 億＝14.7%（Q1 補足資料，Interested-party）；取兩者之間 15%，各年同比率（Assumed）"),
        ("cx_rev", "資本支出 ÷ 營收", 0.20, 0.25, 0.38,
         "FY2025 實績：資本支出 1,061 億 ÷ 營收 4,162 億＝25.5%（stockanalysis 現金流量表，Verified）；公司計畫 2026–2028 年度三年約 5,000 億（EE Times Japan 2026-02-05，Interested-party）÷ 本模型情境 B 三年營收約 1.93 兆＝25.9%",
         "FY2026 單年計畫設備投資 2,100 億 ÷ 5,500 億＝38.2%（前段集中；高值）；FY2029 起擴產計畫結束，固定 25% 可能高估資本支出（Assumed）")):
    IN.cell(r_in, 1, key_); IN.cell(r_in, 2, item_); IN.cell(r_in, 3, "%")
    for j_, v_ in ((4, lo_), (5, base_), (6, hi_)):
        IN.cell(r_in, j_, v_).font = BLUE; IN.cell(r_in, j_).number_format = "0.0%"
    IN.cell(r_in, 5).fill = KEY
    IN.cell(r_in, 7, "Assumed"); IN.cell(r_in, 8, s1_); IN.cell(r_in, 9, s2_)
    wb.defined_names[key_] = DefinedName(key_, attr_text=f"In!$E${r_in}")
    r_in += 1
FCF_YRS = (2026, 2027, 2028, 2029, 2030, 2031, 2032)
fh = PL.max_row + 1; fr = fh + 1
PL.cell(fh, 1, "項目（自由現金流；欄＝年度 × 情境，FY2026 為錨點年、FY2030 起為收斂法）").font = BOLD
PL.cell(fh, 2, "單位")
_col = 3
for s_ in SC:
    for y_ in FCF_YRS:
        PL.cell(fh, _col, f"FY{y_} {s_}").font = BOLD
        R_ = XR[s_]; c_ = XC[y_]
        PL.cell(fr, _col, f"={c_}{R_['ni']}+{c_}{R_['rev']}*da_rev-{c_}{R_['rev']}*cx_rev").number_format = "#,##0.0"
        _col += 1
PL.cell(fh, _col, "說明").font = BOLD
PL.cell(fr, 1, "自由現金流（淨利＋折舊 − 資本支出；不含營運資金變動）"); PL.cell(fr, 2, V2["unit"])
PL.cell(fr, _col, "淨利與營收取自本頁延伸表；折舊＝營收 × In 頁 da_rev、資本支出＝營收 × In 頁 cx_rev（各年同比率，Assumed）；FY2026 若照公司計畫（設備投資 2,100 億、折舊 810 億）約 −45（¥B），固定比率會攤平擴產前段的負值")
_s0 = SO.max_row + 1
for _i, (_a, _b) in enumerate((
        ("[Verified] stockanalysis 現金流量表：2026 年 3 月期（本模型 FY2025）折舊攤銷 66,236、資本支出 106,118、營業現金流 106,407（百萬日圓）", "https://stockanalysis.com/quote/tyo/4062/financials/cash-flow-statement/"),
        ("[Verified] 公司 Q1 補足資料：2027 年 3 月期（本模型 FY2026）設備投資 2,100 億、折舊 810 億（公司計畫，Interested-party）", SUPQ1_URL),
        ("[Verified] EE Times Japan 2026-02-05：2026–2028 年度三年投資約 5,000 億（第一期約 2,200 億，河間 Cell 6，2027 年度起陸續投產）", "https://eetimes.itmedia.co.jp/ee/articles/2602/05/news037.html"),
        ("[Assumed] 折舊 ÷ 營收 15%、資本支出 ÷ 營收 25%（In 頁 da_rev、cx_rev；PL 頁自由現金流列用；FY2029 起擴產計畫結束後可能高估資本支出）", "—"))):
    SO.cell(_s0 + _i, 1, _a); SO.cell(_s0 + _i, 2, _b)
VER_NOTE_EXTRA = "PL 頁最後新增「自由現金流（淨利＋折舊 − 資本支出；不含營運資金變動）」一列（三情境 × FY2026–FY2032，表頭在其上一列），折舊、資本支出＝營收 × In 頁新增的 da_rev（15%）、cx_rev（25%）；不進 EPS 與目標價。"
_t = RD["A1"].value
assert VER_PREV in _t and VER_NOW not in _t, _t
RD["A1"] = _t.replace(VER_PREV, VER_NOW, 1)
for r_ in range(2, RD.max_row + 1):
    _v = RD.cell(r_, 1).value
    if isinstance(_v, str) and _v.startswith("版本："):
        RD.cell(r_, 1, "版本：" + VER_NOW + "（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）。前版：" + _v[3:])
        break
rn = RD.max_row + 1
RD.cell(rn, 1, VER_NOW + "：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁最後新增「上檔空間級距」「投資建議」兩列（第 " + str(GRADE_ROW) + "、" + str(REC_ROW) + " 列，依「上檔空間：加權目標價」的固定規則）；重點觀察見下方三列。" + VER_NOTE_EXTRA)
RD.cell(rn + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"')
for k_, o_ in enumerate(OBS, 1):
    RD.cell(rn + 1 + k_, 1, f"重點觀察 {k_}：" + o_)

order =["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("IBIDEN_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "IMP": IMP, "IMPH": IMPH, "SB": list(SB), "SENS_LAST": SENS_LAST,
               "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "WMT_NOTE": WMT_NOTE, "UNOTE": UNOTE, "QR": QR},
              open(os.environ["IBIDEN_ROWMAP"], "w"), ensure_ascii=False, indent=1)
