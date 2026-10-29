"""分部原則第 3 條：在既有工作表最後追加核對列「檢查：分部合計 − 合併營收」（不位移既有列，數值不變）。

用法（各公司 build_*.py 存檔前呼叫）：
    sys.path.insert(0, os.path.join(<repo 根目錄>, "tools", "segment_check")); from addrow import add_check
    add_check(ws, seg_rows=[5, 8, 9, 10], total_row=4, cols="BCDEF", note="…")
seg_rows：官方分部所在列；total_row：合併營收列；cols：要核對的欄（該欄分部與合併皆有數字或公式才寫）。
回傳核對列列號。
"""
from openpyxl.styles import Font

LABEL = "檢查：分部合計 − 合併營收"


def add_check(ws, seg_rows, total_row, cols, note="", label=LABEL, row=None, fmt="0.000", rnd=6):
    r = row or ws.max_row + 2
    ws.cell(r, 1, label).font = Font(bold=True)
    for c in cols:
        cells = [ws[f"{c}{x}"].value for x in list(seg_rows) + [total_row]]
        if any(v is None or v == "" for v in cells):
            continue
        ws[f"{c}{r}"] = f"=ROUND(" + "+".join(f"{c}{x}" for x in seg_rows) + f"-{c}{total_row},{rnd})"
        ws[f"{c}{r}"].number_format = fmt
    if note:
        ws.cell(r + 1, 1, note)
    return r
