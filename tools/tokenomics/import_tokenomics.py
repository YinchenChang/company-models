#!/usr/bin/env python3
"""Tokenomics 取數工具（各公司共用）：把 Tokenomics 的 IF_／IFW_／IFC_／L1_／SRC_ 具名範圍讀成版本固定的快照 JSON。

用法：
  產生快照：python3 tools/tokenomics/import_tokenomics.py --tokenomics <Tokenomics clone> --names <名稱清單> --out <快照.json> [--optional <清單檔或逗號分隔>]
  重現檢查：python3 tools/tokenomics/import_tokenomics.py --check <快照.json> [--tokenomics <clone>]
            重新讀取快照記錄的同一個 xlsx（model/ 或 model/archive/），任何值差異（相對誤差 > 1e-9）即以代碼 1 結束。
            未給 --tokenomics 時依序找環境變數 TOKENOMICS_DIR、../tokenomics、/home/claude/yinchenchang/tokenomics；
            都找不到時印警告、以代碼 0 結束（略過，不算失敗）。

規則：
  - 只接受下游資料契約第 2 條第 1 項允許的名稱（Tokenomics `docs/plan/Tokenomics_downstream_contract.md`）：
    `IF_`（不含 `IF_Hdr*`）、`IFW_`（四層瀑布，v5.29 起）、`IFC_`（用途欄，v5.29 起）、`L1_`、`SRC_`（只限狀態 Active）；其他名稱報錯。
  - `IFW_*`：與 IF_ 相同，依世代 × 成本情境拆開；取用時須四層一起取（契約第 2 條第 2 項）。
  - `IFC_Conf／IFC_Use／IFC_Layer／IFC_Updated`（Interface R／S／T／U 整欄）：存成 {該列的 IF_／IFW_ 名稱: 值}，
    只保留有具名範圍的列；用來檢查所取名稱的用途限制（例如「上限，不得作預測」）。
  - `SRC_*`（SRC_* 工作表單格 C 欄）：值取 C 欄，另存同列紀錄（指標、低、高、單位、口徑、日期、出處、來源等級、狀態）；狀態不是 Active 即報錯。
  - 名稱清單：一行一個名稱，`#` 之後為註解；行尾加 `optional`（或列在 --optional）＝ Tokenomics 尚無此名稱時記為 {"missing": true} 並警告，不中斷。
  - 讀快取值（openpyxl data_only=True）；需要的格有公式但快取值缺漏時，先以 LibreOffice headless 重算暫存副本再讀。
  - Interface 世代 × 成本情境範圍（C:Q 共 15 欄）依 IF_HdrGen、IF_HdrCost 表頭拆成 {世代: {低成本, 基準, 高成本}}；單格名稱存單值。
    label、unit 取該列 A、B 欄（去掉 A 欄尾端的「[IF_…]」）。L1 名稱（單格 D 欄＝基準）label 取 B 欄、unit 取 G 欄，另存 E／F 欄（低／高）與 H 欄（區間定義）於 range。
"""
import argparse, datetime, json, os, re, shutil, signal, subprocess, sys, tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

GEN_CODES = [('Hopper', 'H100'), ('GB200', 'GB200'), ('GB300', 'GB300'), ('VR200', 'VR200'), ('Rubin Ultra', 'RU')]
REL_TOL = 1e-9
DEFAULT_DIRS = ['../tokenomics', '/home/claude/yinchenchang/tokenomics']


def gen_code(header):
    for key, code in GEN_CODES:
        if str(header).startswith(key):
            return code
    raise SystemExit(f'無法辨識的世代表頭：{header!r}（請更新 GEN_CODES）')


ALLOWED_PREFIX = ('IF_', 'IFW_', 'IFC_', 'L1_', 'SRC_')
SRC_FIELDS = {2: '指標', 4: '低', 5: '高', 6: '單位', 7: '口徑', 9: '日期', 10: '出處', 11: '來源等級', 15: '狀態'}


def check_name(n):
    if n.startswith('IF_Hdr') or not n.startswith(ALLOWED_PREFIX):
        raise SystemExit(f'名稱不合規則：{n}（只接受 IF_（不含 IF_Hdr*）、IFW_、IFC_、L1_、SRC_）')


def read_names(path):
    req, opt = [], []
    for ln in open(path, encoding='utf-8'):
        ln = ln.split('#', 1)[0].strip()
        if not ln:
            continue
        parts = ln.split()
        n = parts[0]
        check_name(n)
        (opt if len(parts) > 1 and parts[1].lower() == 'optional' else req).append(n)
    return req, opt


def parse_ref(text):
    m = re.fullmatch(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?", text)
    if not m:
        raise SystemExit(f'無法解析的具名範圍位置：{text}')
    sh, c1, r1, c2, r2 = m.groups()
    c2, r2 = c2 or c1, r2 or r1
    return sh, column_index_from_string(c1), int(r1), column_index_from_string(c2), int(r2)


def ref_text(sh, c1, r1, c2, r2):
    a = f'{get_column_letter(c1)}{r1}'
    b = f'{get_column_letter(c2)}{r2}'
    return f'{sh}!{a}' if a == b else f'{sh}!{a}:{b}'


def soffice():
    for c in ('soffice', 'libreoffice'):
        if shutil.which(c):
            return c
    raise SystemExit('找不到 soffice（apt-get install -y --no-install-recommends libreoffice-calc-nogui）')


MACRO = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>'''
XLB = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">
<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false">
 <library:element library:name="Module1"/>
</library:library>'''
DLB = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">
<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false"/>'''


def recalc_copy(src, timeout=900):
    """以 LibreOffice headless 重算 src 的暫存副本，回傳副本路徑（原檔不動）。"""
    work = tempfile.mkdtemp(prefix='tk_recalc_')
    dst = os.path.join(work, 'recalc.xlsx')
    shutil.copy(src, dst)
    prof = os.path.join(work, 'profile')
    env_arg = '-env:UserInstallation=' + Path(prof).as_uri()
    env = dict(os.environ, LC_ALL='C.UTF-8', LANG='C.UTF-8')
    subprocess.run([soffice(), env_arg, '--headless', '--norestore', '--terminate_after_init'],
                   capture_output=True, timeout=timeout, env=env)
    std = os.path.join(prof, 'user', 'basic', 'Standard')
    os.makedirs(std, exist_ok=True)
    open(os.path.join(std, 'Module1.xba'), 'w').write(MACRO)
    open(os.path.join(std, 'script.xlb'), 'w').write(XLB)
    open(os.path.join(std, 'dialog.xlb'), 'w').write(DLB)
    mtime = os.path.getmtime(dst)
    pr = subprocess.Popen([soffice(), env_arg, '--headless', '--norestore',
                           'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application',
                           Path(dst).as_uri()], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                          env=env, start_new_session=True)
    try:
        pr.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        os.killpg(pr.pid, signal.SIGKILL)
        pr.communicate()
        raise SystemExit(f'LibreOffice 重算逾時（{timeout} 秒）')
    if os.path.getmtime(dst) == mtime:
        raise SystemExit('LibreOffice 重算失敗（副本未存檔）')
    return dst


def git_head(clone):
    try:
        return subprocess.run(['git', '-C', clone, 'rev-parse', 'HEAD'], capture_output=True, text=True,
                              check=True).stdout.strip()
    except Exception:
        return None


def locate_xlsx(clone, fname):
    for d in ('model', 'model/archive'):
        p = os.path.join(clone, d, fname)
        if os.path.exists(p):
            return p
    return None


def extract(xlsx, names, optional, force_recalc=False):
    """回傳 (items, generations, recalculated)。"""
    wb_f = openpyxl.load_workbook(xlsx, data_only=False)
    dn = wb_f.defined_names
    plan, missing = {}, []
    for n in names:
        if n not in dn:
            if n in optional:
                missing.append(n)
                continue
            raise SystemExit(f'Tokenomics 沒有名稱 {n}（未列為 optional）')
        plan[n] = parse_ref(dn[n].attr_text)
    hdr = {h: parse_ref(dn[h].attr_text) for h in ('IF_HdrGen', 'IF_HdrCost')}

    def cells_of(sh, c1, r1, c2, r2):
        return [(sh, r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    need = [x for p in plan.values() for x in cells_of(*p)]
    wb_v = openpyxl.load_workbook(xlsx, data_only=True)
    gaps = [f'{s}!{get_column_letter(c)}{r}' for s, r, c in need
            if wb_v[s].cell(r, c).value is None and isinstance(wb_f[s].cell(r, c).value, str)
            and str(wb_f[s].cell(r, c).value).startswith('=')]
    recalculated = False
    if gaps or force_recalc:
        if gaps:
            print(f'警告：{len(gaps)} 格快取值缺漏（例 {gaps[:3]}），以 LibreOffice 重算暫存副本後再讀', file=sys.stderr)
        wb_v = openpyxl.load_workbook(recalc_copy(xlsx), data_only=True)
        recalculated = True

    def val(sh, r, c):
        v = wb_v[sh].cell(r, c).value
        return v

    gs, gc1, gr1, gc2, _ = hdr['IF_HdrGen']
    cs, cc1, cr1, _, _ = hdr['IF_HdrCost']
    gen_by_col = {c: wb_v[gs].cell(gr1, c).value for c in range(gc1, gc2 + 1)}
    cost_by_col = {c: wb_v[cs].cell(cr1, c).value for c in range(gc1, gc2 + 1)}
    generations = []
    for c in range(gc1, gc2 + 1):
        g = gen_by_col[c]
        if g not in [x['name'] for x in generations]:
            generations.append({'name': g, 'code': gen_code(g)})

    items = {}
    for n, (sh, c1, r1, c2, r2) in plan.items():
        ws = wb_v[sh]
        it = {'sheet': sh, 'ref': ref_text(sh, c1, r1, c2, r2)}
        if n.startswith('IFC_'):
            if c1 != c2:
                raise SystemExit(f'{n}：IFC_ 名稱預期為單欄，實際為 {it["ref"]}')
            row_name = {}
            for m, d in dn.items():
                if m.startswith(('IF_', 'IFW_')) and not m.startswith('IF_Hdr'):
                    try:
                        msh, _, mr1, _, mr2 = parse_ref(d.attr_text)
                    except SystemExit:
                        continue
                    if msh == sh and mr1 == mr2 and r1 <= mr1 <= r2:
                        row_name.setdefault(mr1, m)
            it['label'] = next((ws.cell(r, c1).value for r in range(r1 - 1, 0, -1)
                                if ws.cell(r, c1).value not in (None, '')), None)
            it['unit'] = None
            it['kind'] = 'by_name'
            it['values'] = {row_name[r]: val(sh, r, c1) for r in sorted(row_name)}
            it['cells'] = {row_name[r]: f'{sh}!{get_column_letter(c1)}{r}' for r in sorted(row_name)}
        elif sh.startswith('SRC_'):
            if not (c1 == c2 and r1 == r2):
                raise SystemExit(f'{n}：SRC_ 名稱預期為單格，實際為 {it["ref"]}')
            rec = {k: val(sh, r1, c) for c, k in SRC_FIELDS.items()}
            if rec['狀態'] != 'Active':
                raise SystemExit(f'{n}：狀態為 {rec["狀態"]!r}，契約只允許取 Active 紀錄')
            it['label'] = rec.pop('指標')
            it['unit'] = rec.pop('單位')
            it['kind'] = 'single'
            it['values'] = val(sh, r1, c1)
            it['cell'] = f'{sh}!{get_column_letter(c1)}{r1}'
            it['record'] = rec
        elif sh == 'L1':
            if not (c1 == c2 and r1 == r2):
                raise SystemExit(f'{n}：L1 名稱預期為單格，實際為 {it["ref"]}')
            it['label'] = ws.cell(r1, 2).value
            it['unit'] = ws.cell(r1, 7).value
            it['condition'] = ws.cell(r1, 3).value
            it['kind'] = 'single'
            it['values'] = val(sh, r1, c1)
            it['cell'] = f'{sh}!{get_column_letter(c1)}{r1}'
            it['range'] = {'低': val(sh, r1, 5), '高': val(sh, r1, 6), 'def': ws.cell(r1, 8).value,
                           'cells': {'低': f'{sh}!E{r1}', '高': f'{sh}!F{r1}'}}
            ext = {'srcId': ws.cell(r1, 12).value, '外部低': ws.cell(r1, 13).value, '外部高': ws.cell(r1, 14).value}
            if ext['srcId'] and ext['srcId'] != '—':
                it['external'] = ext
        else:
            lab = ws.cell(r1, 1).value
            it['label'] = re.sub(r'\s*\[(IF|IFW|L1)_[^\]]*\]\s*$', '', str(lab)).strip() if lab is not None else None
            it['unit'] = ws.cell(r1, 2).value
            if r1 == r2 and c1 == c2:
                it['kind'] = 'single'
                it['values'] = val(sh, r1, c1)
                it['cell'] = f'{sh}!{get_column_letter(c1)}{r1}'
            elif r1 == r2 and (c1, c2) == (gc1, gc2) and sh == gs:
                it['kind'] = 'gen_cost'
                v, cells = {}, {}
                for c in range(c1, c2 + 1):
                    g, k = gen_by_col[c], cost_by_col[c]
                    v.setdefault(g, {})[k] = val(sh, r1, c)
                    cells.setdefault(g, {})[k] = f'{sh}!{get_column_letter(c)}{r1}'
                it['values'], it['cells'] = v, cells
            else:
                it['kind'] = 'vector'
                it['values'] = [val(*x) for x in cells_of(sh, c1, r1, c2, r2)]
                it['cells'] = [f'{s}!{get_column_letter(c)}{r}' for s, r, c in cells_of(sh, c1, r1, c2, r2)]
        items[n] = it
    for n in missing:
        print(f'警告：Tokenomics 尚無名稱 {n}（optional），快照記為 missing', file=sys.stderr)
        items[n] = {'missing': True}
    ordered = {n: items[n] for n in names}
    return ordered, generations, recalculated


def find_clone(arg):
    cands = [arg] if arg else [os.environ.get('TOKENOMICS_DIR')] + DEFAULT_DIRS
    for c in cands:
        if c and os.path.exists(os.path.join(c, 'model', 'CURRENT')):
            return os.path.abspath(c)
    return None


def same(a, b):
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(same(a[k], b[k]) for k in a)
    if isinstance(a, list) and isinstance(b, list):
        return len(a) == len(b) and all(same(x, y) for x, y in zip(a, b))
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return a == b or abs(a - b) <= REL_TOL * max(abs(a), abs(b))
    return a == b


def diff_items(old, new, path=''):
    out = []
    if isinstance(old, dict) and isinstance(new, dict):
        for k in sorted(set(old) | set(new), key=str):
            if k not in old or k not in new:
                out.append(f'{path}{k}：{"快照沒有" if k not in old else "重新讀取沒有"}')
            else:
                out += diff_items(old[k], new[k], f'{path}{k}.')
    elif not same(old, new):
        out.append(f'{path.rstrip(".")}：快照 {old!r} ≠ 重新讀取 {new!r}')
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--tokenomics')
    ap.add_argument('--names')
    ap.add_argument('--out')
    ap.add_argument('--optional', help='optional 名稱：清單檔路徑或逗號分隔')
    ap.add_argument('--check', metavar='SNAPSHOT')
    ap.add_argument('--force-recalc', action='store_true', help='一律先以 LibreOffice 重算副本再讀（測試用）')
    a = ap.parse_args()

    if a.check:
        snap = json.load(open(a.check, encoding='utf-8'))
        clone = find_clone(a.tokenomics)
        if not clone:
            print('警告：找不到 Tokenomics clone（--tokenomics／TOKENOMICS_DIR），略過 --check', file=sys.stderr)
            return 0
        src = snap['source']
        xlsx = locate_xlsx(clone, src['file'])
        if not xlsx:
            print(f'失敗：Tokenomics clone {clone} 找不到快照的檔案 {src["file"]}（model/ 或 model/archive/）')
            return 1
        head = git_head(clone)
        if head and head != src.get('commit'):
            print(f'注意：clone HEAD {head[:7]} ≠ 快照 commit {str(src.get("commit"))[:7]}；以同一檔名 {src["file"]} 比對')
        names = list(snap['items'])
        optional = set(src.get('optional', []))
        items, gens, _ = extract(xlsx, names, optional | {n for n, v in snap['items'].items() if v.get('missing')},
                                 a.force_recalc)
        d = diff_items(snap['items'], items) + diff_items({'generations': src.get('generations')},
                                                          {'generations': gens})
        if d:
            print(f'--check 失敗：{len(d)} 項差異')
            for x in d[:30]:
                print('  ', x)
            return 1
        nm = sum(1 for v in items.values() if v.get('missing'))
        print(f'--check 通過：{len(items)} 個名稱（missing {nm}）與快照一致（{src["file"]}，相對誤差 ≤ {REL_TOL}）')
        return 0

    if not (a.tokenomics and a.names and a.out):
        ap.error('產生快照需要 --tokenomics、--names、--out（或改用 --check）')
    clone = os.path.abspath(a.tokenomics)
    cur = open(os.path.join(clone, 'model', 'CURRENT'), encoding='utf-8').read().strip()
    xlsx = os.path.join(clone, 'model', cur)
    req, opt = read_names(a.names)
    if a.optional:
        extra = (read_names(a.optional)[0] + read_names(a.optional)[1]) if os.path.exists(a.optional) \
            else [x.strip() for x in a.optional.split(',') if x.strip()]
        for n in extra:
            check_name(n)
        opt += [n for n in extra if n not in opt]
    names = req + [n for n in opt if n not in req]
    items, gens, recalculated = extract(xlsx, names, set(opt), a.force_recalc)
    m = re.search(r'_v(\d+\.\d+)\.xlsx$', cur)
    snap = {'source': {'file': cur, 'current': cur, 'version': 'v' + m.group(1) if m else None,
                       'commit': git_head(clone), 'extractedAt': datetime.date.today().isoformat(),
                       'repo': 'YinchenChang/Tokenomics', 'recalculated': recalculated,
                       'optional': [n for n in names if n in opt], 'generations': gens,
                       'note': '主值取「基準」成本情境；低成本／高成本只作敏感度。每 GW＝IT 關鍵電力（Tokenomics Interface 第 1 列）。'},
            'items': items}
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, 'w', encoding='utf-8') as f:
        json.dump(snap, f, ensure_ascii=False, indent=1)
        f.write('\n')
    nm = sum(1 for v in items.values() if v.get('missing'))
    print(f'快照：{a.out}（{cur}，commit {str(snap["source"]["commit"])[:7]}；{len(items)} 個名稱，missing {nm}）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
