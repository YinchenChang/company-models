"""來源複核（API 版，2026-10-28 起；說明見 tools/source_audit/README.md）。

只產出複核結果，不改任何模型；不符者由各模型的維護串依工作單處理。

對象：標記為 Verified 或 Interested-party、基準值為數字的輸入列。
  1. Siliconomics 快照列（siliconomics/reader.py 的 ROWS）
  2. 各公司模型 dist 最新 xlsx 的 In 工作表
  來源欄寫「同上」者沿用上一列來源；來源為 Siliconomics／markets／Tokenomics 的繼承列不重複查（在源頭查）。
做法：同一來源的列合併成一組（每組至多 8 列），以 Batch API 送 Sonnet，附 web_search 與 web_fetch，
  請模型打開來源核對數字。打不開的頁面一律標「來源打不開」，不得改用快取、鏡像或其他繞道。
費用：依 usage 逐波累計（Batch 價：輸入 1、輸出 5 美元／百萬 tokens；搜尋 0.01 美元／次），
  下一波預估會超過 MAX_USD 就停。已完成的組記在 out/state.jsonl，重跑時跳過（可分次跑完）。
輸出：tools/source_audit/out/source_audit.xlsx 與 .md。
"""
import glob
import json
import os
import re
import sys
import time
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
STATE = os.path.join(OUT, "state.jsonl")
API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
MODEL = os.environ.get("AUDIT_MODEL", "claude-sonnet-5-5")
MAX_USD = float(os.environ.get("MAX_USD", "40"))
MAX_GROUPS = int(os.environ.get("MAX_GROUPS", "0") or 0)      # 0 = 不限（試跑時設小數字）
WAVE = int(os.environ.get("WAVE", "60"))
POLL_MAX_MIN = int(os.environ.get("POLL_MAX_MIN", "120"))
PRICE_IN, PRICE_OUT, PRICE_SEARCH = 1.0 / 1e6, 5.0 / 1e6, 0.01   # Batch 價
TAGS = ("Verified", "Interested-party")
INHERITED = re.compile(r"^\s*(Siliconomics|markets/|Tokenomics|SIF_|TIF_)")
GROUP_MAX = 8


# ---------- 收集 ----------
def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def collect():
    rows = []
    sys.path.insert(0, os.path.join(ROOT, "siliconomics"))
    import reader as S  # noqa: E402
    last = ""
    for r in S.ROWS:
        if "id" not in r:
            continue
        src = (r.get("src") or "").strip()
        if src.startswith("同上") or not src:
            src = last + (src[2:] if src.startswith("同上") else "")
        else:
            last = src
        if any(t in str(r.get("tag")) for t in TAGS) and _num(r.get("base")):
            rows.append(dict(scope="siliconomics", file="siliconomics 快照", key=r["id"], label=r["label"], unit=r.get("unit") or "",
                             lo=r.get("lo"), base=r["base"], hi=r.get("hi"), tag=r["tag"], src=src))
    import openpyxl
    for x in sorted(glob.glob(f"{ROOT}/*/dist/*.xlsx")):
        d = os.path.relpath(x, ROOT).split("/")[0]
        if d == "siliconomics":
            continue
        files = sorted(glob.glob(f"{ROOT}/{d}/dist/*.xlsx"))
        if x != files[-1]:
            continue
        try:
            ws = openpyxl.load_workbook(x, read_only=True, data_only=True)["In"]
        except Exception:
            continue
        last = ""
        for r in ws.iter_rows(min_row=4, values_only=True):
            if not r or len(r) < 8 or not r[0]:
                continue
            src = str(r[7] or "").strip()
            if src.startswith("同上") or not src:
                src = last + (src[2:] if src.startswith("同上") else "")
            else:
                last = src
            if any(t in str(r[6] or "") for t in TAGS) and _num(r[4]) and src and not INHERITED.match(src):
                rows.append(dict(scope=d, file=os.path.basename(x), key=str(r[0]), label=str(r[1] or ""), unit=str(r[2] or ""),
                                 lo=r[3], base=r[4], hi=r[5], tag=str(r[6]), src=src))
    return rows


def groups(rows):
    by = {}
    for r in rows:
        by.setdefault((r["scope"], r["src"]), []).append(r)
    out = []
    for (scope, src), rs in by.items():
        for i in range(0, len(rs), GROUP_MAX):
            part = rs[i:i + GROUP_MAX]
            gid = f"{scope}|{part[0]['key']}|{len(part)}"
            out.append(dict(gid=gid, scope=scope, src=src, rows=part))
    # 先 Siliconomics，再各公司輪流取一組（額度用完時每家都有覆蓋）
    sil = [g for g in out if g["scope"] == "siliconomics"]
    per = {}
    for g in out:
        if g["scope"] != "siliconomics":
            per.setdefault(g["scope"], []).append(g)
    rr = []
    k = 0
    while any(k < len(v) for v in per.values()):
        rr += [v[k] for v in per.values() if k < len(v)]
        k += 1
    return sil + rr


# ---------- API ----------
SYS = """你是審計助理，任務是核對投資模型輸入值與其標示來源是否相符。只做核對，不提出新估計。
規則：
1. 依「來源／說明」欄找原始來源：有網址就用 web_fetch 打開；沒有網址就用 web_search 找該機構的原始發布（新聞稿、財報、法說、研究機構公開摘要），再打開。
2. 只用合法公開途徑。頁面被封鎖、需登入、付費牆、robots 拒絕或回傳錯誤時，不得改用快取、鏡像、轉載站或其他繞道，直接判「來源打不開」。
3. 單位為 % 的模型值以小數表示（0.155 = 15.5%）。單位換算（十億／百萬、% 與小數、季與年、匯率）若來源欄已寫明換算方式，依其換算後再比；Derived 換算步驟合理即可。
4. 判定只能是以下五種之一：
   相符：來源數字與模型值一致（容許四捨五入，相對差 ≤2%）
   部分相符：方向與量級一致但數字不同、口徑不同、或只找到其中一部分
   不符：來源明確寫出不同數字
   來源打不開：找到來源但無法讀取
   找不到來源：搜尋後找不到該來源或該數字
5. 最後只輸出一個 JSON 陣列，不要其他文字：
[{"key":"<鍵>","verdict":"<五種之一>","source_value":"<來源上的原文數字與單位，沒有就空字串>","url":"<實際核對的網址>","note":"<一句中文說明，≤60字>"}]"""


def prompt(g):
    lines = [f"模型：{g['scope']}（{g['rows'][0]['file']}）", f"來源／說明：{g['src']}", "", "待核對的列："]
    for r in g["rows"]:
        rng = f"；區間 {r['lo']}～{r['hi']}" if _num(r["lo"]) and _num(r["hi"]) else ""
        lines.append(f"- 鍵 {r['key']}｜{r['label']}｜基準 {r['base']} {r['unit']}{rng}｜標記 {r['tag']}")
    return "\n".join(lines)


TOOLS = [{"type": "web_search_20250305", "name": "web_search", "max_uses": 3},
         {"type": "web_fetch_20250910", "name": "web_fetch", "max_uses": 4, "max_content_tokens": 20000}]


def api(path, body=None, method="GET"):
    req = urllib.request.Request(f"https://api.anthropic.com/v1/{path}",
                                 data=json.dumps(body).encode() if body is not None else None, method=method,
                                 headers={"x-api-key": API_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"})
    with urllib.request.urlopen(req, timeout=300) as r:
        return r.read().decode()


def cost(usage):
    u = usage or {}
    tin = u.get("input_tokens", 0) + 1.25 * u.get("cache_creation_input_tokens", 0) + 0.1 * u.get("cache_read_input_tokens", 0)
    s = (u.get("server_tool_use") or {}).get("web_search_requests", 0)
    return tin * PRICE_IN + u.get("output_tokens", 0) * PRICE_OUT + s * PRICE_SEARCH, s, (u.get("server_tool_use") or {}).get("web_fetch_requests", 0)


def parse(text):
    m = re.search(r"\[\s*\{.*\}\s*\]", text, re.S)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except Exception:
        return None


def run_wave(wave):
    ids = {f"g{i:04d}": g for i, g in enumerate(wave)}
    reqs = [{"custom_id": k, "params": {"model": MODEL, "max_tokens": 6000, "system": SYS, "tools": TOOLS,
                                       "messages": [{"role": "user", "content": prompt(g)}]}} for k, g in ids.items()]
    b = json.loads(api("messages/batches", {"requests": reqs}, "POST"))
    print("batch", b["id"], len(reqs), flush=True)
    t0 = time.time()
    while b["processing_status"] != "ended" and time.time() - t0 < POLL_MAX_MIN * 60:
        time.sleep(30)
        b = json.loads(api(f"messages/batches/{b['id']}"))
    if b["processing_status"] != "ended":
        api(f"messages/batches/{b['id']}/cancel", {}, "POST")
        print("逾時，已取消", flush=True)
        return [], 0.0
    out, spent = [], 0.0
    for line in api(f"messages/batches/{b['id']}/results").splitlines():
        if not line.strip():
            continue
        it = json.loads(line)
        g = ids[it["custom_id"]]
        rec = dict(gid=g["gid"], scope=g["scope"], ok=False, results=[], cost=0.0, searches=0, fetches=0, stop="")
        if it["result"]["type"] == "succeeded":
            msg = it["result"]["message"]
            c, s, f = cost(msg.get("usage"))
            text = "".join(x.get("text", "") for x in msg["content"] if x.get("type") == "text")
            res = parse(text)
            rec.update(cost=c, searches=s, fetches=f, stop=msg.get("stop_reason", ""), ok=res is not None, results=res or [], raw=text[-600:])
            spent += c
        else:
            rec["stop"] = it["result"]["type"]
        out.append(rec)
    return out, spent


# ---------- 主程式 ----------
def load_state():
    done, spent = {}, 0.0
    if os.path.exists(STATE):
        for line in open(STATE, encoding="utf-8"):
            r = json.loads(line)
            spent += r.get("cost", 0)
            if r.get("ok"):
                done[r["gid"]] = r
    return done, spent


def main():
    os.makedirs(OUT, exist_ok=True)
    rows = collect()
    gs = groups(rows)
    done, spent_before = load_state()
    todo = [g for g in gs if g["gid"] not in done]
    if MAX_GROUPS:
        todo = todo[:MAX_GROUPS]
    print(f"列 {len(rows)}｜組 {len(gs)}｜已完成 {len(done)}｜本次待做 {len(todo)}｜先前花費 {spent_before:.2f}｜上限 {MAX_USD}", flush=True)
    spent, per_group = spent_before, None
    i = 0
    while i < len(todo) and API_KEY:
        n = WAVE if per_group else min(WAVE, 20)
        est = (per_group or 0.08) * n
        if spent + est > MAX_USD:
            n = int((MAX_USD - spent) / (per_group or 0.08))
            if n < 1:
                print(f"已達上限：累計 {spent:.2f}", flush=True)
                break
        wave = todo[i:i + n]
        recs, c = run_wave(wave)
        spent += c
        with open(STATE, "a", encoding="utf-8") as f:
            for r in recs:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        okc = [r["cost"] for r in recs if r["cost"]]
        if okc:
            per_group = max(sum(okc) / len(okc), 0.01) * 1.15
        print(f"波次 {i}-{i + len(wave)}：本波 {c:.2f}｜累計 {spent:.2f}｜每組約 {per_group or 0:.3f}", flush=True)
        i += len(wave)
    report(rows, gs, spent)


VERDICTS = ["不符", "部分相符", "來源打不開", "找不到來源", "相符"]


def report(rows, gs, spent):
    done, _ = load_state()
    allrecs = [json.loads(l) for l in open(STATE, encoding="utf-8")] if os.path.exists(STATE) else []
    searches = sum(r.get("searches", 0) for r in allrecs)
    fetches = sum(r.get("fetches", 0) for r in allrecs)
    gmap = {g["gid"]: g for g in gs}
    out = []
    for gid, rec in done.items():
        g = gmap.get(gid)
        if not g:
            continue
        res = {str(x.get("key")): x for x in rec["results"] if isinstance(x, dict)}
        for r in g["rows"]:
            x = res.get(r["key"], {})
            v = x.get("verdict", "")
            out.append(dict(r, verdict=v if v in VERDICTS else (v or "模型未回覆此列"), source_value=x.get("source_value", ""),
                            url=x.get("url", ""), note=x.get("note", "")))
    order = {v: i for i, v in enumerate(VERDICTS)}
    out.sort(key=lambda r: (order.get(r["verdict"], -1), r["scope"], r["key"]))
    from collections import Counter
    cnt = Counter(r["verdict"] for r in out)
    total = len(rows)
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "摘要"
    summ = [("項目", "數值"), ("應核對列數（Verified／Interested-party，非繼承）", total), ("已核對列數", len(out)),
            ("未核對列數（額度用完或尚未跑）", total - len(out)), ("累計花費（美元，估）", round(spent, 2)), ("上限（美元）", MAX_USD),
            ("網頁搜尋次數", searches), ("網頁擷取次數", fetches), ("模型", MODEL)]
    summ += [(f"判定：{v}", cnt.get(v, 0)) for v in VERDICTS] + [("判定：模型未回覆此列", cnt.get("模型未回覆此列", 0))]
    for s in summ:
        ws.append(s)
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 46
    d = wb.create_sheet("明細")
    hdr = ["判定", "模型", "檔案", "鍵", "項目", "單位", "低", "基準", "高", "標記", "模型來源欄", "來源上的數字", "核對網址", "說明"]
    d.append(hdr)
    fill = {"不符": "F8CBAD", "部分相符": "FFE699", "來源打不開": "D9D9D9", "找不到來源": "D9D9D9", "相符": "C6E0B4"}
    for r in out:
        d.append([r["verdict"], r["scope"], r["file"], r["key"], r["label"], r["unit"], r["lo"], r["base"], r["hi"], r["tag"],
                  r["src"], r["source_value"], r["url"], r["note"]])
        if r["verdict"] in fill:
            d.cell(d.max_row, 1).fill = PatternFill("solid", fgColor=fill[r["verdict"]])
    for c, w in zip("ABCDEFGHIJKLMN", [10, 14, 30, 16, 36, 10, 8, 10, 8, 14, 60, 28, 40, 50]):
        d.column_dimensions[c].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = d.dimensions
    wb.save(os.path.join(OUT, "source_audit.xlsx"))
    md = ["# 來源複核結果", "", f"- 應核對 {total} 列，已核對 {len(out)} 列；累計花費約 {spent:.2f} 美元（上限 {MAX_USD}）；搜尋 {searches} 次、擷取 {fetches} 次。", ""]
    md += [f"- {v}：{cnt.get(v, 0)}" for v in VERDICTS]
    md += ["", "## 不符與部分相符", "", "| 判定 | 模型 | 鍵 | 項目 | 基準 | 來源上的數字 | 說明 |", "|---|---|---|---|---|---|---|"]
    for r in out:
        if r["verdict"] in ("不符", "部分相符"):
            md.append(f"| {r['verdict']} | {r['scope']} | {r['key']} | {r['label']} | {r['base']} {r['unit']} | {r['source_value']} | {r['note']} |")
    open(os.path.join(OUT, "source_audit.md"), "w", encoding="utf-8").write("\n".join(md) + "\n")
    print("\n".join(md[:10]))


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--dry":
        rs = collect()
        gs = groups(rs)
        from collections import Counter
        print(len(rs), "列", len(gs), "組", Counter(r["scope"] for r in rs).most_common(8))
        print(prompt(gs[0]))
    else:
        main()
