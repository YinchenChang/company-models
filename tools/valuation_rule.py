# 估值規則區塊（Val 頁）：把「十年倍數歷史」與「依循環位置取分位的規則值」放進公司 xlsx，
# 讓 Excel 成為規則值的唯一事實來源（總覽 extract.py 與 HTML 展示層都讀這裡）。
# 規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底倍數依循環位置取分位：見頂年 25 分位、中段中位數、谷底 75 分位。
# 多個來源時，各來源分別取分位再平均，四捨五入到兩位；min／max 取各來源的極值。
# 用法：tools/memory_model/build.py 與 tsmc/build_tsmc.py 在 Val 頁底部呼叫 add_rule_block。
from openpyxl.styles import Font, PatternFill

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")
YEARS = list(range(2016, 2026))
COLS = "BCDEFG"


def _hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h)
        c.font = BOLD
        c.fill = HDR


def add_sec_table(ws, r0, title, headers, rows, note=None):
    """輸入表（藍字）：rows 為 [(年度, 值…)]；回傳 (下一列, {年度: 列號})。"""
    ws.cell(r0, 1, title).font = BOLD
    _hdr(ws, r0 + 1, headers)
    where = {}
    for i, row in enumerate(rows):
        r = r0 + 2 + i
        ws.cell(r, 1, row[0])
        for j, v in enumerate(row[1:], 2):
            c = ws.cell(r, j, v)
            c.font = BLUE
        where[row[0]] = r
    r = r0 + 2 + len(rows)
    if note:
        ws.cell(r, 1, note)
        r += 1
    return r + 1, where


def add_rule_block(ws, r0, series, fmtp, intro):
    """series：[{id, name, kind, sources:[{label, vals(10 個：數值、公式字串或 None), tag, note, use}], cases:[{case, pos, q, basis_label, basis}]}]
    use=False 的來源只列統計（參考／核對），不進「採用」；ref_rows=True 的來源另列「參考分位」列。
    值若為含 {r} 的公式字串，{r} 代入該年度所在列。回傳下一個空列。"""
    ws.cell(r0, 1, "6. 十年倍數與規則值（2016–2025 年底；見頂年 25 分位、中段中位數、谷底 75 分位）").font = BOLD
    ws.cell(r0 + 1, 1, intro)
    r = r0 + 3
    for s in series:
        sid = s["id"]
        srcs = s["sources"]
        ws.cell(r, 1, f"{s['name']}：十年歷史倍數").font = BOLD
        _hdr(ws, r + 1, ["年底"] + [x["label"] for x in srcs] + ["標記／說明"])
        y0 = r + 2
        for i, y in enumerate(YEARS):
            ws.cell(y0 + i, 1, str(y))
            for j, x in enumerate(srcs):
                v = x["vals"][i]
                if v is not None:
                    if isinstance(v, str) and "{r}" in v:
                        v = v.replace("{r}", str(y0 + i))
                    c = ws.cell(y0 + i, 2 + j, v)
                    if not (isinstance(v, str) and v.startswith("=")):
                        c.font = BLUE
                    c.number_format = "0.00"
        notes = [x.get("note") for x in srcs if x.get("note")]
        ws.cell(y0, 2 + len(srcs), "；".join(f"{x['label']}：{x['tag']}" for x in srcs))
        for k, t in enumerate(notes):
            ws.cell(y0 + 1 + k, 2 + len(srcs), t)
        y9 = y0 + len(YEARS) - 1
        r = y9 + 1
        st = {}
        for lab, fn in (("25 分位", "PERCENTILE({rng},0.25)"), ("中位數", "MEDIAN({rng})"), ("75 分位", "PERCENTILE({rng},0.75)"),
                        ("最小", "MIN({rng})"), ("最大", "MAX({rng})")):
            ws.cell(r, 1, lab + "（各來源）")
            for j in range(len(srcs)):
                col = COLS[j]
                ws.cell(r, 2 + j, "=" + fn.format(rng=f"{col}{y0}:{col}{y9}")).number_format = "0.00"
            st[lab] = r
            r += 1
        use = [COLS[j] for j, x in enumerate(srcs) if x.get("use", True)]
        ref = [COLS[j] for j, x in enumerate(srcs) if x.get("ref_rows")]
        adopt = {}
        for lab, agg in (("25 分位", "AVERAGE"), ("中位數", "AVERAGE"), ("75 分位", "AVERAGE"), ("最小", "MIN"), ("最大", "MAX")):
            ws.cell(r, 1, f"採用 {lab}（{sid}）").font = BOLD
            cells = ",".join(f"{c}{st[lab]}" for c in use)
            ws.cell(r, 2, f"=ROUND({agg}({cells}),2)").number_format = "0.00"
            adopt[lab] = r
            r += 1
        for lab in ("25 分位", "中位數", "75 分位"):
            if ref:
                ws.cell(r, 1, f"參考 {lab}（{sid}；不採用的來源）")
                ws.cell(r, 2, f"=ROUND({ref[0]}{st[lab]},2)").number_format = "0.00"
                r += 1
        r += 1
        ws.cell(r, 1, f"{s['name']}：規則值").font = BOLD
        _hdr(ws, r + 1, ["情形", "循環位置", "倍數", "基準", "每股基準", "隱含股價", "相對目前", "現價 ÷ 基準"])
        r += 2
        for cs in s["cases"]:
            q = {25: "25 分位", 50: "中位數", 75: "75 分位"}[cs["q"]]
            ws.cell(r, 1, f"規則值：{cs['case']}（{sid}）")
            ws.cell(r, 2, cs["pos"])
            ws.cell(r, 3, f"=B{adopt[q]}").number_format = "0.00"
            ws.cell(r, 4, cs["basis_label"])
            ws.cell(r, 5, cs["basis"]).number_format = fmtp
            ws.cell(r, 6, f"=C{r}*E{r}").number_format = "#,##0.0"
            ws.cell(r, 7, f"=F{r}/px-1").number_format = "0.0%"
            ws.cell(r, 8, f"=px/E{r}").number_format = "0.00"
            r += 1
        r += 2
    return r
