# 相對倍數調整（Comparable 的目標倍數＝同業前瞻倍數中位數 × 本公司 10 年相對同業的倍數比）
# 依據：Andy 2026-10-27 15:20「照建議做」（AI半導體 目標價總表 v0.2.3 評等偏向分析的建議 1）。
# 狀態（2026-10-29）：資料已換成 Bloomberg 前瞻倍數（Andy 10-29 取回）；**仍未套用到任何模型**（下一步：接入 54 個建檔程式）。
#   10-27 以年底落後 P/E 試算，15 檔評等翻轉中多數是谷底年度雜訊造成的偽訊號，因此不採用（目標價總表 v0.2.4「相對倍數試算」頁）。
# 問題：Comparable 直接套同業中位數，等於假設本公司長期的折價／溢價（韓國折價、A 股地緣折價、
#       代工組裝低毛利、獨占溢價等）會在 12 個月內消失；評等因此主要反映「現價倍數高低」。
# 做法：
#   本公司 10 年中位數 ÷ 同業各自 10 年中位數的中位數 ＝ 原始比值；夾在 [RM_LO, RM_HI]；
#   資料不足（本公司有效年數 < MIN_YEARS，或有效同業 < MIN_PEERS）時係數＝1（不調整）。
#   Comparable 目標價（相對倍數調整後）＝ Comparable 目標價（市場主流同業法）× 係數；加權目標價改用調整後的值。
#   「Comparable 目標價（市場主流同業法）」與「上檔空間：Comparable 目標價」維持不調整（Andy 2026-10-25：看與市場主流的差距）。
# 資料：tools/data/fwd_mult_hist.csv——Bloomberg BDH 月底前瞻倍數（P/E：BEST_PE_RATIO；記憶體 P/B：BEST_PX_BPS_RATIO；BF＝未來 12 個月），
#   每個年度窗口（前一年 11 月～當年 10 月）取有效月值中位數（有效月數 ≥ 6），只存這 10 個年度值，不存月資料。
#   月資料原檔在 Drive 03-AI半導體/選股與資料/。舊落後代理 tools/data/rel_mult_hist.csv 保留供對照。
# 用法（各公司建檔程式，在 Val 頁所有內容寫完之後、存檔之前）：
#   import relative_multiple as RM
#   RM.add_block(wb, V, own="2382.TW", peers=[...], metric="PE", comp_ref=f"Val!$B${T['comptp']}", pxfmt="#,##0")
#   然後把加權目標價公式中的 Comparable 改成名稱 tp_comp_adj。
import csv
import os
import statistics
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data", "fwd_mult_hist.csv")   # 2026-10-29 起：Bloomberg 前瞻倍數；舊落後代理 rel_mult_hist.csv 保留供對照
YEARS = list(range(2017, 2027))   # 年度窗口：前一年 11 月～當年 10 月（2026＝2025-11～2026-10）
RM_LO, RM_HI = 0.5, 2.0          # [Assumed] 夾值上下限；Andy 可改（改這裡即全部模型一致）
MIN_YEARS, MIN_PEERS = 7, 2       # [Assumed] 資料門檻（7 年：排除只有 2021–2025 的序列，避免 AI 高峰期偏誤）
VALID = {"PE": (0.0, 200.0), "PB": (0.0, 20.0)}   # 有效值區間（不含下限）；沿用十年 P/E 規則「負值或 >200 剔除」

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")


def norm(t):
    t = str(t).strip()
    if "／" in t:
        t = t.split("／")[1]
    return t


def load(path=DATA):
    """回傳 {(ticker, metric): {year: (value, url)}}；value 為 float。"""
    out = {}
    with open(path, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if not r["year"] or r["value"] in ("", None):
                out.setdefault((r["ticker"], r["metric"]), {})
                continue
            out.setdefault((r["ticker"], r["metric"]), {})[int(r["year"])] = (float(r["value"]), r["source_url"])
    return out


def _valid(v, metric):
    lo, hi = VALID[metric]
    return lo < v <= hi


def summary(own, peers, metric="PE", hist=None):
    """Python 端同口徑計算（供總表、稽核與測試用；Excel 端以公式重算，兩者須一致）。"""
    hist = hist if hist is not None else load()

    def med(t):
        vals = [v for y, (v, _) in hist.get((norm(t), metric), {}).items() if y in YEARS and _valid(v, metric)]
        return (statistics.median(vals) if len(vals) >= MIN_YEARS else None), len(vals)
    om, on = med(own)
    pm = []
    for _, t in peers:
        m, _n = med(t)
        if m is not None:
            pm.append(m)
    pmed = statistics.median(pm) if pm else None
    ok = om is not None and len(pm) >= MIN_PEERS
    raw = om / pmed if ok else None
    fac = min(max(raw, RM_LO), RM_HI) if ok else 1.0
    return {"own_med": om, "own_n": on, "peer_meds": pm, "peer_med": pmed, "n_peers": len(pm), "raw": raw, "factor": fac, "ok": ok}


def add_block(wb, ws, own, peers, metric="PE", comp_ref=None, pxfmt="#,##0", own_name="本公司", r0=None, hist=None):
    """在 ws（Val 頁）最後寫入「相對倍數調整」區塊；定義名稱 rm_factor、tp_comp_adj。
    peers：[(名稱, 代號)]，與 Comparable 同業表相同的公司（順序不限）。
    comp_ref：Comparable 目標價（市場主流同業法）所在儲存格，例如 "Val!$B$130"。回傳 dict（各列列號）。"""
    hist = hist if hist is not None else load()
    if r0 is None:
        r0 = ws.max_row + 3
    lab = "P/E" if metric == "PE" else "P/B"
    ws.cell(r0, 1, f"0.5 相對倍數調整（Comparable 目標倍數＝同業中位數 × 本公司 10 年相對同業的 {lab} 比；Andy 2026-10-27 新增，放在本頁最後以免位移既有儲存格）").font = BOLD
    ws.cell(r0 + 1, 1, f"資料：{YEARS[0]}–{YEARS[-1]} 各年度窗口（前一年 11 月～當年 10 月）Bloomberg 月底前瞻 {lab}（BF）的中位數（tools/data/fwd_mult_hist.csv），[Verified] 來源、[Derived] 彙總；"
                       f"有效值 {VALID[metric][0]:g} < 值 ≤ {VALID[metric][1]:g}，無效值以「剔除:」文字列示、不進中位數。"
                       f"本公司或同業有效年數 < {MIN_YEARS} 不納入；有效同業 < {MIN_PEERS} 家或本公司不足時係數＝1（不調整）。")
    h = r0 + 2
    heads = ["公司", "代號", "角色"] + [f"{y - 1}/11–{y}/10" for y in YEARS] + ["有效年數", "10 年中位數", "納入值", "來源"]
    for j, x in enumerate(heads, 1):
        c = ws.cell(h, j, x)
        c.font = BOLD
        c.fill = HDR
    CY = {y: 4 + i for i, y in enumerate(YEARS)}          # D..M
    CN, CMED, CIN, CSRC = 14, 15, 16, 17                  # N O P Q
    from openpyxl.utils import get_column_letter as L
    rows = [(own_name, own, "本公司")] + [(n, t, "同業") for n, t in peers]
    r = h + 1
    first = r
    for nm, tk, role in rows:
        tk = norm(tk)
        ws.cell(r, 1, nm)
        ws.cell(r, 2, tk)
        ws.cell(r, 3, role)
        d = hist.get((tk, metric), {})
        urls = []
        for y in YEARS:
            if y in d:
                v, u = d[y]
                if u and u not in urls:
                    urls.append(u)
                if _valid(v, metric):
                    c = ws.cell(r, CY[y], v)
                    c.font = BLUE
                    c.number_format = "0.00"
                else:
                    ws.cell(r, CY[y], f"剔除:{v:g}")
        rng = f"{L(CY[YEARS[0]])}{r}:{L(CY[YEARS[-1]])}{r}"
        ws.cell(r, CN, f"=COUNT({rng})")
        ws.cell(r, CMED, f'=IF({L(CN)}{r}>0,MEDIAN({rng}),"")').number_format = "0.00"
        ws.cell(r, CIN, f'=IF({L(CN)}{r}>={MIN_YEARS},{L(CMED)}{r},"")').number_format = "0.00"
        ws.cell(r, CSRC, "；".join(urls) if urls else ("未找到" if not d else "未找到有效值"))
        r += 1
    last = r - 1
    pin = f"{L(CIN)}{first + 1}:{L(CIN)}{last}" if last > first else None
    r += 1
    R = {}

    def row(key, label, formula, fmt, note=""):
        nonlocal r
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, formula)
        c.number_format = fmt
        if note:
            ws.cell(r, 5, note)
        R[key] = r
        r += 1
    row("own", "本公司 10 年中位數（納入值）", f"={L(CIN)}{first}", "0.00", "空白＝本公司有效年數不足")
    row("npeer", "納入的同業家數", f"=COUNT({pin})" if pin else "=0", "0")
    row("pmed", "同業 10 年中位數的中位數", f'=IF(B{r - 1}>0,MEDIAN({pin}),"")' if pin else '=""', "0.00")
    row("raw", "原始比值（本公司 ÷ 同業）", f'=IF(AND(ISNUMBER(B{R["own"]}),B{R["npeer"]}>={MIN_PEERS}),B{R["own"]}/B{R["pmed"]},"")', "0.00",
        "＜1＝本公司長期比同業便宜（折價），＞1＝長期溢價")
    row("lo", "夾值下限", RM_LO, "0.00", "[Assumed] 避免單一歷史異常主導；tools/relative_multiple.py 的 RM_LO")
    ws.cell(R["lo"], 2).font = BLUE
    row("hi", "夾值上限", RM_HI, "0.00", "[Assumed] tools/relative_multiple.py 的 RM_HI")
    ws.cell(R["hi"], 2).font = BLUE
    row("fac", "相對倍數調整係數", f'=IF(ISNUMBER(B{R["raw"]}),MIN(MAX(B{R["raw"]},B{R["lo"]}),B{R["hi"]}),1)', "0.00",
        "資料不足時＝1（不調整）")
    ws.cell(R["fac"], 1).font = BOLD
    ws.cell(R["fac"], 2).fill = KEY
    if comp_ref:
        row("adj", "Comparable 目標價（相對倍數調整後）", f"={comp_ref}*B{R['fac']}", pxfmt,
            "＝Comparable 目標價（市場主流同業法）× 調整係數；加權目標價使用本列")
        ws.cell(R["adj"], 1).font = BOLD
        ws.cell(R["adj"], 2).fill = KEY
        wb.defined_names["tp_comp_adj"] = DefinedName("tp_comp_adj", attr_text=f"{ws.title}!$B${R['adj']}")
    wb.defined_names["rm_factor"] = DefinedName("rm_factor", attr_text=f"{ws.title}!$B${R['fac']}")
    R["head"] = h
    R["first"] = first
    R["last"] = last
    R["title"] = r0
    return R


# ---- 一次接入（2026-10-29）：各建檔程式在 wb.save 之前呼叫 RM.apply(wb, "<資料夾>") ----
MEMORY_PB = ("skhynix", "samsung", "micron", "kioxia", "sndk")   # 記憶體用前瞻 P/B，其餘前瞻 P/E
_TS = os.path.join(HERE, "target_summary")


def _universe():
    with open(os.path.join(_TS, "universe.csv"), encoding="utf-8") as f:
        return {r["folder"]: r for r in csv.DictReader(f)}


def _peers(folder):
    with open(os.path.join(_TS, "peers.csv"), encoding="utf-8") as f:
        return [(r["name"], r["ticker"]) for r in csv.DictReader(f) if r["folder"] == folder]


def apply(wb, folder):
    """把相對倍數調整接入一個公司模型：
    (1) Val 頁最後加「0.5 相對倍數調整」區塊（同業取 tools/target_summary/peers.csv，本公司代號取 universe.csv）；
    (2) 加權目標價公式中的 Comparable 改讀 tp_comp_adj；
    (3) Sens 頁中「＝Comparable 目標價」的引用格改讀 tp_comp_adj（敏感度與加權目標價同口徑）。
    「Comparable 目標價（市場主流同業法）」「上檔空間：Comparable 目標價」、Recon 的 Comparable 對帳維持不調整（Andy 2026-10-25）。"""
    import re
    V = wb["Val"]
    tpw = comp = None
    for r in range(1, V.max_row + 1):
        a, b = str(V.cell(r, 1).value or ""), V.cell(r, 2).value
        if a.startswith("目標價（DCF 與 Comparable 加權）") and isinstance(b, str) and "tp_w_dcf" in b:
            m = re.search(r"\(1-tp_w_dcf\)\*\$?B\$?(\d+)", b)
            assert m, (folder, b)
            tpw, comp = r, int(m.group(1))
            break
    assert tpw, f"{folder}：Val 頁找不到加權目標價公式"
    u = _universe()[folder]
    peers = _peers(folder)
    assert peers, f"{folder}：peers.csv 無同業"
    fmt = V.cell(tpw, 2).number_format or "#,##0"
    R = add_block(wb, V, own=u["own_ticker"], peers=peers, metric="PB" if folder in MEMORY_PB else "PE",
                  comp_ref=f"Val!$B${comp}", pxfmt=fmt, own_name=u["公司"])
    V.cell(tpw, 2).value = re.sub(r"\(1-tp_w_dcf\)\*\$?B\$?\d+", "(1-tp_w_dcf)*tp_comp_adj", V.cell(tpw, 2).value)
    rx = re.compile(rf"^=Val!\$?B\$?{comp}$")                                   # 「Comparable 目標價」欄直接引用
    rx_w = re.compile(rf"(\(1-[^()]*\)\*)Val!\$?B\$?{comp}(?!\d)")                  # 公式內的（1 − 權重）× Comparable
    n_sens = 0
    if "Sens" in wb.sheetnames:
        S = wb["Sens"]
        for row in S.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                if rx.match(c.value):
                    c.value = "=tp_comp_adj"
                    n_sens += 1
                    for up in range(1, 12):
                        h = S.cell(c.row - up, c.column) if c.row - up >= 1 else None
                        if h is not None and h.value == "Comparable 目標價":
                            h.value = "Comparable 目標價（相對倍數調整後）"
                            break
                elif rx_w.search(c.value):
                    c.value = rx_w.sub(r"\1tp_comp_adj", c.value)
                    n_sens += 1
    R.update(tpw=tpw, comp=comp, n_sens=n_sens)
    if "README" in wb.sheetnames:
        _obs1(wb, V, tpw, comp, fmt)
        ws = wb["README"]
        ws.cell(ws.max_row + 2, 1,
                f"2026-10-29：加權目標價中的 Comparable 改用相對倍數調整後的值（Val 第 0.5 節：同業前瞻倍數中位數 × 本公司 10 年相對同業的倍數比，"
                f"Bloomberg 前瞻倍數；Andy 2026-10-27「照建議做」、10-29 資料取回）。「Comparable 目標價（市場主流同業法）」維持不調整。重點觀察 1 的第一句已改為公式；重點觀察 2、3 與其他說明中的目標價與上檔空間數字為調整前撰寫，以 Val 頁為準。")
    return R


def _xl_text(t):
    """長文字切成 ≤ 200 字的字串常數，以 & 串接（Excel 公式內字串常數上限 255 字）。"""
    t = t.replace('"', '""')
    return "&".join(f'"{t[i:i + 200]}"' for i in range(0, len(t), 200)) or '""'


def _obs1(wb, V, tpw, comp, fmt):
    """README「重點觀察 1」第一句（目標價、上檔空間、級距）改為公式，數字隨重算更新；原說明保留並標明為調整前撰寫。"""
    import re
    ws = wb["README"]
    cell = None
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(("重點觀察 1：", "1. 加權目標價")):
                cell = c
                break
        if cell:
            break
    if cell is None:
        return
    lab = {}
    for r in range(1, V.max_row + 1):
        a = str(V.cell(r, 1).value or "")
        for k, p in (("dcf", "DCF 目標價（情境加權）"), ("up", "上檔空間：加權目標價"), ("upc", "上檔空間：Comparable 目標價"),
                     ("tier", "上檔空間級距"), ("rec", "投資建議")):
            if k not in lab and a.startswith(p) and V.cell(r, 2).value is not None:
                lab[k] = r
    if len(lab) < 5:
        return
    t = cell.value
    head = "重點觀察 1：" if t.startswith("重點觀察 1：") else "1. "
    if head == "1. ":
        t = "重點觀察 1：" + t[len("1. 加權"):]
    m = re.match(r"重點觀察 1：目標價\s*([^\d\s（]*)\s*[\d,\.]+", t)
    cur = m.group(1) if m else ""
    d = re.search(r"現價\s*[\d,\.]+（([^）]+)）", t)
    pxd = f"（{d.group(1)}）" if d else ""
    i = t.find("Comparable 單獨為")
    j = t.find("。", i) if i >= 0 else -1
    if j < 0 and head == "1. ":          # 健策型：第一個逗號段之後的說明保留
        j = t.find("），") + 1 if "），" in t[:80] else t.find("，", t.find("上檔空間"))
    rest = t[j + 1:].strip() if j >= 0 else t[len("重點觀察 1："):]
    nf = fmt if fmt and fmt != "General" else "#,##0"
    pct = '"0.0%"'
    sp = '"+0.0%;−0.0%"'
    B = lambda k: f"Val!$B${lab[k]}"
    pre = f"(tp_w_dcf*{B('dcf')}+(1-tp_w_dcf)*Val!$B${comp})/px-1"
    parts = [_xl_text(f"{head}目標價 {cur}"), f'TEXT(Val!$B${tpw},"{nf}")',
             _xl_text("（DCF "), f'TEXT({B("dcf")},"{nf}")', _xl_text("、Comparable（相對倍數調整後）"), f'TEXT(tp_comp_adj,"{nf}")',
             _xl_text("；DCF 權重 "), 'TEXT(tp_w_dcf,"0%")', _xl_text("），比現價 "), f'TEXT(px,"{nf}")', _xl_text(pxd or " "),
             f'IF({B("up")}>=0,"高 "&TEXT({B("up")},{pct}),"低 "&TEXT(-{B("up")},{pct}))',
             _xl_text("，級距 "), B("tier"), _xl_text("、投資建議 "), B("rec"),
             _xl_text("。相對倍數調整係數 "), 'TEXT(rm_factor,"0.00")',
             _xl_text("（本公司 10 年前瞻倍數中位數 ÷ 同業的中位數，夾在 0.5–2.0；Val 第 0.5 節），調整前加權上檔空間 "),
             f"TEXT({pre},{sp})", _xl_text("；Comparable 單獨（市場主流同業法、未調整）為 "), f"TEXT({B('upc')},{sp})", _xl_text("。")]
    if rest:
        parts.append(_xl_text("以下為 2026-10-29 相對倍數調整前撰寫的說明，其中數字為調整前：" + rest))
    f = "=" + "&".join(parts)
    cell.value = f
    for row in ws.iter_rows():      # README 開頭若有寫死的投資建議行，改為公式
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("投資建議："):
                c.value = (f'="投資建議："&{B("rec")}&"（依上檔空間固定規則；加權目標價 {cur}"&TEXT(Val!$B${tpw},"{nf}")'
                           f'&"、上檔空間 "&TEXT({B("up")},{sp})&"、級距 "&{B("tier")}&"；見 Val 第 0 節與第 0.5 節）"')
