# 個股模型 In 頁版面核對（全 repo 共用；固定規格見 docs/plan/個股模型原則_命題加營收分部.md 第 8 節；
# Andy 2026-10-29 14:26 同意、15:03 指示套用到其他適用的個股）。
#
# 適用：In 頁有「低／基準／高」三欄、且有情境專屬列的模型。情境專屬列的判定（is_scn）：
#   (a) 鍵名含一個獨立的情境字母 A／B／C（前一字元為底線、小寫或數字或在開頭，後一字元為底線、數字或結尾），
#       且把該字母換成另一個情境字母後的鍵也存在（同族兄弟，例如 tp_wA／tp_wB、cqA_x／cqB_x、dA／dC、pace_A／pace_C）；或
#   (b) 項目文字以「情境 A：」或「情境 A（…）：」開頭，或含「（情境 A）」「情境權重 A」（只有單一情境的列，例如 gm_adj_C28）。
#   只在其他文字中提到情境（例如「收斂法情境 B 的 FY2032 營收年增」「台達電模型，情境 A」）不算。
#   (b) 的例外（2026-10-30 放寬甲，AI半導體 請求）：該列 B 欄以後任一格含「非本模型情境驅動」（例如上游模型情境 A 的參考值），
#   不算情境專屬。(a) 同族兄弟不受此例外影響。
# 規格：
#   1. In!D3／F3 欄名為「敏感度低」「敏感度高」（只用於 Sens 頁，不進目標價）。
#   2. In 頁最後有一節，A 欄文字含「情境驅動對照」：下一列表頭 D／E／F＝情境 A／B／C；其下每列一個情境驅動，三欄並排，
#      每格是 =E<列> 公式，引用 H 欄「引用 In 頁 k1／k2／k3；…」所列鍵的列。某情境無專屬輸入時可填常數 0：
#      情境 B 不需說明；情境 A 或 C 填 0 時，H 欄須寫「情境 A 無專屬輸入列」或「情境 C 無專屬輸入列」（2026-10-30 放寬甲）。
#      D／E／F 只能是 =E<列> 或數字常數；引用列若可判定情境字母（同族兄弟或標籤「情境 X」），須與所在欄（D＝A、E＝B、F＝C）一致。
#   3. 上方所有情境專屬列都必須被這一節引用（不能漏列）；情境權重列鍵為 scn_w，合計＝1。
#   4. 本節純展示：其他工作表與 In 頁上方的公式、定義名稱都不得引用本節。
#   5.（警告，不判失敗）Sens、Recon、In 以外工作表的公式不應直接讀 In 頁 D／F 欄（敏感度低／高），含指向 D／F 的定義名稱（以實際指向判定，不看名稱字尾）：
#      情境值應拆成情境專屬列（例如 g27_A／g27_B／g27_C）並列入對照；共識區間、舊版對照等非主模型用途可保留，於模型 README 說明。
# 用法：
#   python3 tools/in_layout/check_in_layout.py <成品 xlsx>      # 單檔；不符即失敗（結束碼 1）
#   python3 tools/in_layout/check_in_layout.py --all [--md 報告.md]  # 掃描全部 */dist/*.xlsx，只報告（結束碼 0）
import glob, os, re, sys
import openpyxl

EXEMPT = "非本模型情境驅動"
LAB_PAT = re.compile(r"^情境 ?[ABC][：:]|^情境 ?[ABC]（[^）]*）[：:]|（情境 ?[ABC]）|情境權重 ?[ABC](?![A-Za-z])")
KEY_OK = re.compile(r"^[A-Za-z0-9_]+$")


def scn_letter(key, label, keys):
    """情境專屬列所屬的情境字母（A／B／C）；判斷不出時回傳 None。先看鍵名同族兄弟，再看標籤「情境 X」。"""
    if KEY_OK.match(key):
        for i, ch in enumerate(key):
            if ch not in "ABC":
                continue
            prev_ok = i == 0 or key[i - 1] == "_" or key[i - 1].islower() or key[i - 1].isdigit()
            next_ok = i == len(key) - 1 or key[i + 1] == "_" or key[i + 1].isdigit()
            if prev_ok and next_ok and any(key[:i] + o + key[i + 1:] in keys for o in "ABC" if o != ch):
                return ch
    m = re.search(r"情境 ?([ABC])", label) if LAB_PAT.search(label) else None
    return m.group(1) if m else None


def is_scn(key, label, keys, exempt=False):
    if KEY_OK.match(key):
        for i, ch in enumerate(key):
            if ch not in "ABC":
                continue
            prev_ok = i == 0 or key[i - 1] == "_" or key[i - 1].islower() or key[i - 1].isdigit()
            next_ok = i == len(key) - 1 or key[i + 1] == "_" or key[i + 1].isdigit()
            if prev_ok and next_ok and any(key[:i] + o + key[i + 1:] in keys for o in "ABC" if o != ch):
                return True
    return bool(LAB_PAT.search(label)) and not exempt


WARN_SKIP = {"In", "Sens", "Recon"}


def lohi_uses(wbf, last):
    """Sens／Recon／In 以外讀敏感度低／高欄的公式格（第 5 項警告）：
    定義名稱以「實際指向 In!$D$n 或 In!$F$n（n < 情境驅動對照）」認定，不看名稱字尾；另加直接寫 In!D／F 的公式。"""
    names = set()
    for n, dn_ in wbf.defined_names.items():
        for t, ref in dn_.destinations:
            m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", ref)
            if t == "In" and m and m.group(1) in ("D", "F") and int(m.group(2)) < last:
                names.add(n)
    pat_n = re.compile(r"(?<![A-Za-z0-9_.])(" + "|".join(sorted(map(re.escape, names), key=len, reverse=True)) + r")(?![A-Za-z0-9_(])", re.I) if names else None
    hits = []
    for sh in wbf.worksheets:
        if sh.title in WARN_SKIP:
            continue
        for row in sh.iter_rows():
            for c in row:
                f = c.value
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                f2 = re.sub(r'"[^"]*"', "", f)  # 不看字串常數
                if (pat_n and pat_n.search(f2)) or any(int(x) < last for x in re.findall(r"In!\$?[DF]\$?(\d+)(?![\d:])", f2)):
                    hits.append(f"{sh.title}!{c.coordinate}")
    return hits


def check(path):
    """回傳 (適用, 失敗清單, 情境專屬列數, 警告清單)"""
    wbv = openpyxl.load_workbook(path, data_only=True)
    if "In" not in wbv.sheetnames:
        return False, [], 0, []
    ws = wbv["In"]
    wf = openpyxl.load_workbook(path, data_only=False)["In"]
    hdr = [str(ws.cell(3, c).value or "") for c in (4, 5, 6)]
    if hdr[1] != "基準":
        return False, [], 0, []
    sec = [r for r in range(4, ws.max_row + 1) if "情境驅動對照" in str(ws.cell(r, 1).value or "")]
    last = (sec[0] if sec else ws.max_row + 1)
    keyrow = {str(wf.cell(r, 1).value): r for r in range(4, last) if wf.cell(r, 1).value}
    def exempt(r):
        return any(EXEMPT in str(wf.cell(r, c).value or "") for c in range(2, min(wf.max_column, 12) + 1))
    scn_rows = {r for r in range(4, last) if wf.cell(r, 1).value and is_scn(str(wf.cell(r, 1).value), str(ws.cell(r, 2).value or ""), keyrow, exempt(r))}
    if not scn_rows:
        return False, [], 0, []
    fail = []
    lh = lohi_uses(wf.parent, last)
    warn = [f"{len(lh)} 格讀敏感度低／高欄（Sens／Recon 以外）：" + "、".join(lh[:6]) + ("…" if len(lh) > 6 else "")] if lh else []
    if hdr[0] != "敏感度低" or hdr[2] != "敏感度高":
        fail.append(f"In!D3／F3 應為「敏感度低」「敏感度高」：{hdr[0]}／{hdr[2]}")
    if len(sec) != 1:
        fail.append(f"「情境驅動對照」一節應恰好一個：{len(sec)}")
        return True, fail, len(scn_rows), warn
    h = sec[0] + 1
    if [str(ws.cell(h, c).value or "")[:4] for c in (4, 5, 6)] != ["情境 A", "情境 B", "情境 C"]:
        fail.append(f"情境驅動對照表頭 D／E／F 應為情境 A／B／C：{[ws.cell(h, c).value for c in (4, 5, 6)]}")
    rows, refd = {}, set()
    r = h + 1
    while ws.cell(r, 1).value:
        k = ws.cell(r, 1).value
        v = [ws.cell(r, c).value for c in (4, 5, 6)]
        rows[k] = v
        if not all(isinstance(x, (int, float)) for x in v):
            fail.append(f"情境驅動 {k} 三欄應皆為數值：{v}")
        hnote = str(wf.cell(r, 8).value or "")
        m = re.match(r"引用 In 頁 ([^；]+)", hnote)
        keys = m.group(1).split("／") if m else []
        fs = [wf.cell(r, c).value for c in (4, 5, 6)]
        refs = [int(f[2:]) for f in fs if isinstance(f, str) and re.fullmatch(r"=E\d+", f)]
        consts = [j for j, f in enumerate(fs) if not isinstance(f, str)]
        # 2026-10-30 補：只接受 =E<列> 或數字常數（=0 之類公式不放行）；引用列的情境字母須與所在欄一致
        others = [j for j, f in enumerate(fs) if isinstance(f, str) and not re.fullmatch(r"=E\d+", f)]
        if others:
            fail.append(f"情境驅動 {k} 的 {'、'.join('DEF'[j] for j in others)} 欄只能是 =E<列> 或數字常數：{[fs[j] for j in others]}")
        for j, f in enumerate(fs):
            if isinstance(f, str) and re.fullmatch(r"=E\d+", f):
                rr = int(f[2:])
                lt = scn_letter(str(wf.cell(rr, 1).value or ""), str(ws.cell(rr, 2).value or ""), keyrow) if 4 <= rr < last else None
                if lt and lt != "ABC"[j]:
                    fail.append(f"情境驅動 {k} 的 {'DEF'[j]} 欄（情境 {'ABC'[j]}）引用了情境 {lt} 的列 {wf.cell(rr, 1).value}")
        def const_ok(j):
            if fs[j] != 0:
                return False
            return j == 1 or f"情境 {'ABC'[j]} 無專屬輸入列" in hnote
        if [keyrow.get(x) for x in keys] != refs or not all(const_ok(j) for j in consts):
            fail.append(f"情境驅動 {k} 的公式 {fs} 與 H 欄所列鍵 {keys} 不符")
        refd.update(refs)
        r += 1
    if "scn_w" not in rows:
        fail.append("缺情境權重列 scn_w")
    elif all(isinstance(x, (int, float)) for x in rows["scn_w"]) and abs(sum(rows["scn_w"]) - 1) > 1e-9:
        fail.append(f"情境權重合計應為 1：{rows['scn_w']}")
    # 4. 本節純展示：不得被引用
    wbf = wf.parent
    sec_rows = set(range(sec[0], ws.max_row + 1))
    bad = []
    for sh in wbf.worksheets:
        for row in sh.iter_rows():
            for c in row:
                f = c.value
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                if sh.title == "In" and c.row in sec_rows:
                    continue
                rows_ = [int(x) for x in re.findall(r"In!\$?[A-Z]{1,3}\$?(\d+)", f)]
                if sh.title == "In":
                    rows_ += [int(x) for x in re.findall(r"(?<![A-Za-z!_$])\$?[A-Z]{1,3}\$?(\d+)(?![\d(])", f)]
                if any(x in sec_rows for x in rows_):
                    bad.append(f"{sh.title}!{c.coordinate}")
    for name, dn_ in wbf.defined_names.items():
        for t, ref in dn_.destinations:
            if t == "In" and any(int(x) in sec_rows for x in re.findall(r"\$?[A-Z]{1,3}\$?(\d+)", ref)):
                bad.append(f"定義名稱 {name}")
    if bad:
        fail.append("情境驅動對照被引用（應為純展示）：" + "、".join(bad[:8]))
    miss = sorted(scn_rows - refd)
    if miss:
        fail.append("情境專屬列未列入對照：" + "、".join(f"{wf.cell(x, 1).value}（第 {x} 列）" for x in miss))
    return True, fail, len(scn_rows), warn


def main(argv):
    if argv and argv[0] == "--all":
        md = argv[argv.index("--md") + 1] if "--md" in argv else None
        root = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
        res = []
        for d in sorted(glob.glob(os.path.join(root, "*", "dist"))):
            xs = sorted(x for x in glob.glob(os.path.join(d, "*.xlsx")))
            if not xs:
                continue
            x = xs[-1]
            try:
                app, fail, n, warn = check(x)
            except Exception as e:  # 讀檔失敗也列出
                app, fail, n, warn = True, [f"讀檔失敗：{e}"], 0, []
            co = os.path.basename(os.path.dirname(d))
            if app:
                nw = int(re.match(r"(\d+)", warn[0]).group(1)) if warn else 0
                res.append((co, os.path.basename(x), n, "符合" if not fail else "不符", nw, "；".join(fail)[:200]))
        lines = ["| 資料夾 | 成品 | 情境專屬列 | 結果 | 讀低／高欄（警告） | 說明 |", "|---|---|---|---|---|---|"] + [f"| `{a}` | {b} | {c} | {d} | {w} | {e} |" for a, b, c, d, w, e in res]
        print("\n".join(lines))
        print(f"\n適用 {len(res)} 檔；符合 {sum(1 for x in res if x[3] == '符合')}、不符 {sum(1 for x in res if x[3] != '符合')}；有第 5 項警告 {sum(1 for x in res if x[4])} 檔")
        if md:
            open(md, "w", encoding="utf-8").write("\n".join(lines) + "\n")
        return 0
    app, fail, n, warn = check(argv[0])
    for w in warn:
        print("警告（第 5 項，不判失敗）：" + w)
    if not app:
        print("In 頁版面核對：不適用（無 In 頁三欄或無情境專屬輸入列）")
        return 0
    if fail:
        print("In 頁版面核對失敗：")
        for f in fail:
            print("  -", f)
        return 1
    print(f"In 頁版面核對通過（情境專屬 {n} 列皆已列入對照）")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
