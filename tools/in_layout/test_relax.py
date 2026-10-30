# check_in_layout.py 放寬甲（2026-10-30）的合成測試：python3 tools/in_layout/test_relax.py
import os, sys, tempfile
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from check_in_layout import check


def make(rows, sec):
    """rows: [(key, label, note)]；sec: [(key, D, E, F, H)]"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "In"
    ws.cell(3, 4, "敏感度低"); ws.cell(3, 5, "基準"); ws.cell(3, 6, "敏感度高")
    r = 4
    for k, lab, note in rows:
        ws.cell(r, 1, k); ws.cell(r, 2, lab); ws.cell(r, 5, 0.1); ws.cell(r, 8, note); r += 1
    rows_at = {ws.cell(i, 1).value: i for i in range(4, r)}
    r += 1; ws.cell(r, 1, "J. 情境驅動對照"); r += 1
    ws.cell(r, 4, "情境 A"); ws.cell(r, 5, "情境 B"); ws.cell(r, 6, "情境 C"); r += 1
    for k, d, e, f, h in sec:
        ws.cell(r, 1, k)
        for c, v in zip((4, 5, 6), (d, e, f)):
            ws.cell(r, c, v if not isinstance(v, str) or v.startswith("=") else "=E%d" % rows_at[v])
        ws.cell(r, 8, h); r += 1
    p = tempfile.mktemp(suffix=".xlsx"); wb.save(p)
    return p


def run(name, rows, sec, expect_ok, expect_n=None):
    p = make(rows, sec)
    # openpyxl 寫的公式沒有快取值，「三欄應皆為數值」必然觸發，不屬本測試範圍，略過
    app, fail, n, _ = check(p)
    fail = [f for f in fail if "三欄應皆為數值" not in f]
    ok = app and not fail
    res = "OK" if ok == expect_ok and (expect_n is None or n == expect_n) else "NG"
    print(f"{res} {name}: app={app} n={n} fail={fail}")
    return res == "OK"


W = ("scn_w", "w_A", "w_B", "w_C", "引用 In 頁 w_A／w_B／w_C")
base = [("w_A", "情境權重 A", ""), ("w_B", "情境權重 B", ""), ("w_C", "情境權重 C", "")]
cases = [
    ("只有情境 C 列、A／B 填 0 且 H 註明", base + [("gm_adj_C28", "情境 C：毛利率調整", "")],
     [W, ("gm_adj", 0, 0, "gm_adj_C28", "引用 In 頁 gm_adj_C28；情境 A 無專屬輸入列")], True),
    ("只有情境 C 列、A 填 0 未註明 → 失敗", base + [("gm_adj_C28", "情境 C：毛利率調整", "")],
     [W, ("gm_adj", 0, 0, "gm_adj_C28", "引用 In 頁 gm_adj_C28")], False),
    ("A 填非 0 常數 → 失敗", base + [("gm_adj_C28", "情境 C：毛利率調整", "")],
     [W, ("gm_adj", 0.5, 0, "gm_adj_C28", "引用 In 頁 gm_adj_C28；情境 A 無專屬輸入列")], False),
    ("上游情境 A 參考值標註豁免 → 不算情境專屬、不必列入", base + [("inl_26", "上游營收（情境 A）", "上游模型參考值，非本模型情境驅動")],
     [W], True, 3),
    ("上游情境 A 參考值未標註 → 仍須列入（漏列失敗）", base + [("inl_26", "上游營收（情境 A）", "上游模型參考值")],
     [W], False),
    ("同族兄弟不受豁免影響", base + [("g_A", "成長", "非本模型情境驅動"), ("g_C", "成長", "")],
     [W], False),
]
# 2026-10-30 補舊缺口：=0 公式不放行；情境字母須與所在欄一致
cases += [
    ("D 欄寫 =0 公式 → 失敗", base + [("gm_adj_C28", "情境 C：毛利率調整", "")],
     [W, ("gm_adj", "=0", 0, "gm_adj_C28", "引用 In 頁 gm_adj_C28；情境 A 無專屬輸入列")], False),
    ("情境 C 列放在 D 欄（情境 A）→ 失敗", base + [("gm_adj_C28", "情境 C：毛利率調整", "")],
     [W, ("gm_adj", "gm_adj_C28", 0, 0, "引用 In 頁 gm_adj_C28；情境 C 無專屬輸入列")], False),
    ("同族兄弟 A／C 對調欄位 → 失敗", base + [("g_A", "成長", ""), ("g_C", "成長", "")],
     [W, ("g", "g_C", 0, "g_A", "引用 In 頁 g_C／g_A")], False),
    ("同族兄弟正確欄位 → 通過", base + [("g_A", "成長", ""), ("g_C", "成長", "")],
     [W, ("g", "g_A", 0, "g_C", "引用 In 頁 g_A／g_C")], True),
]
ok = all(run(*c) for c in cases)
sys.exit(0 if ok else 1)
