# tools/sens_coverage/check_sens_coverage.py 的單元測試（暫存小活頁簿）。
# 執行：python3 -m pytest -q tools/sens_coverage/test_check_sens_coverage.py  或  python3 tools/sens_coverage/test_check_sens_coverage.py
import os, subprocess, sys, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import check_sens_coverage as C  # noqa: E402

HE = [("he_h", "IF_HoldEcon｜H100", 6.8, 9.7, 16.7), ("he_b2", "IF_HoldEcon｜GB200", 6.6, 8.8, 14.3),
      ("he_b3", "IF_HoldEcon｜GB300", 8.4, 10.9, 18.5), ("he_vr", "IF_HoldEcon｜VR200", 9.2, 12.2, 19.8)]


def make(path, in_rows, sens_rows, a2=None, in_sheet=True, three_col=True, sens=True):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if in_sheet:
        ws = wb.create_sheet("In")
        ws.append(["In"]); ws.append([None])
        ws.append(["鍵", "項目", "單位"] + (["敏感度低", "基準", "敏感度高"] if three_col else ["值"]) + ["標記", "來源", "說明"])
        for k, lab, d, e, f, *note in in_rows:
            ws.append([k, lab, "x", d, e, f, "Judgment", "src", note[0] if note else None])
        ws.append([None]); ws.append(["情境驅動對照（純展示）"]); ws.append(["鍵", "項目", "單位", "情境 A", "情境 B", "情境 C"])
        ws.append(["zz_A", "不該被讀", "x", 1, 2, 3])  # 本節以下不列入
    if sens:
        s = wb.create_sheet("Sens")
        s.append(["Sens — 測試"]); s.append([a2])
        s.append(["驅動", "設定", "方向", "加權目標價", "相對基準", "鍵", "說明"])
        for drv, key, note in sens_rows:
            s.append([drv, "低／高", "低", 100, 0, key, note])
    wb.save(path)
    return path


BASE_IN = [("tax", "有效稅率", 0.13, 0.16, 0.19), ("same", "低高同基準", 1, 1, 1), ("only_hi", "只有高", None, 0, 1)] + HE
BASE_SENS = [("有效稅率：低", "tax", "重算值"), ("只有高：高", "only_hi", "重算值"), ("IF_HoldEcon 四世代同時：低", "he_all", "重算值")]


def tmp(name):
    return os.path.join(tempfile.mkdtemp(prefix="senscov_"), name)


def test_pass_with_he_all():
    r = C.check(make(tmp("ok.xlsx"), BASE_IN, BASE_SENS))
    assert r["applicable"] and not r["missing"] and r["keys"] == 6 and "IF_HoldEcon" in r["group_cov"]


def test_missing_he_fails_even_if_exempted():
    sens = [x for x in BASE_SENS if x[1] != "he_all"]
    a2 = "he_h、he_b2、he_b3、he_vr 不列：測試用豁免（群組不得豁免）。"
    p = make(tmp("nohe.xlsx"), BASE_IN, sens, a2=a2)
    r = C.check(p)
    assert sorted(m[0] for m in r["missing"]) == ["he_b2", "he_b3", "he_h", "he_vr"]
    assert all("不得豁免" in m[5] for m in r["missing"])
    assert C.main([p]) == 1


def test_single_he_member_case_not_enough():
    sens = [x for x in BASE_SENS if x[1] != "he_all"] + [("GB300 持有成本：高", "he_b3", "重算值")]
    r = C.check(make(tmp("he1.xlsx"), BASE_IN, sens))
    assert {m[0] for m in r["missing"]} == {"he_h", "he_b2", "he_b3", "he_vr"}


def test_he_group_by_text_only_fails():
    # 文字「IF_HoldEcon 四世代同時」但鍵欄不是 he_all、也沒寫出四個成員鍵：不算（核對反例 3 的前提）
    sens = [x for x in BASE_SENS if x[1] != "he_all"] + [("Tokenomics IF_HoldEcon 四世代同時：低", "", "")]
    r = C.check(make(tmp("hetxt.xlsx"), BASE_IN, sens))
    assert {m[0] for m in r["missing"]} == {"he_h", "he_b2", "he_b3", "he_vr"}


def test_he_group_by_all_member_keys_in_row():
    sens = [x for x in BASE_SENS if x[1] != "he_all"] + [("四世代同取低", "he_h／he_b2／he_b3／he_vr", "")]
    assert not C.check(make(tmp("he4.xlsx"), BASE_IN, sens))["missing"]


def test_exemption_in_a2_passes():
    ins = BASE_IN + [("inc_m", "增量 EBITDA 率", 0.75, 0.85, 0.95)]
    a2 = "相對基準＝…。inc_m（增量 EBITDA 率）不列：A、C 對稱，對加權目標價結構上恆為 0。"
    p = make(tmp("ex.xlsx"), ins, BASE_SENS, a2=a2)
    r = C.check(p)
    assert not r["missing"] and r["exempt"][0][0] == "inc_m"
    assert C.main([p]) == 0


def test_exemption_in_note_column_passes():
    ins = BASE_IN + [("gl_rl", "RL 營收長期年增", -0.05, 0.05, 0.15)]
    sens = BASE_SENS + [("（無）", "", "gl_rl 豁免：重算變動小於 0.001%")]
    assert not C.check(make(tmp("ex2.xlsx"), ins, sens))["missing"]


def test_missing_plain_key_fails_and_mention_without_exempt_word_is_not_exempt():
    ins = BASE_IN + [("wacc_adj", "WACC 調整", -0.01, 0, 0.01)]
    p = make(tmp("miss.xlsx"), ins, BASE_SENS, a2="wacc_adj 見 Val 頁")  # 提到鍵但沒有豁免字
    r = C.check(p)
    assert [m[0] for m in r["missing"]] == ["wacc_adj"]


def test_key_mentioned_in_case_text_counts():
    ins = BASE_IN + [("cip_ai", "在建工程中 AI 占比", 0.8, 0.9, 1)]
    sens = BASE_SENS + [("在建工程不另增 MW（cip_ai＝0）：低", "other_key", "")]
    assert not C.check(make(tmp("txt.xlsx"), ins, sens))["missing"]


def test_not_model_note_skipped():
    ins = BASE_IN + [("cons_tp", "市場共識目標價", 75, 245, 800, "低／高＝最低／最高（只作對照，不進模型）")]
    r = C.check(make(tmp("nm.xlsx"), ins, BASE_SENS))
    assert not r["missing"] and r["not_model"] == ["cons_tp"]


def test_no_in_sheet_not_applicable():
    p = make(tmp("noin.xlsx"), [], BASE_SENS, in_sheet=False)
    assert C.check(p)["applicable"] is False
    assert C.main([p]) == 0


def test_in_without_three_cols_not_applicable():
    p = make(tmp("old.xlsx"), [("x", "舊模板", 1, 2, 3)], [], three_col=False)
    assert C.check(p)["applicable"] is False


def test_applicable_without_sens_fails():
    p = make(tmp("nosens.xlsx"), BASE_IN, [], sens=False)
    r = C.check(p)
    assert r["error"] and len(r["missing"]) == 6
    assert C.main([p]) == 1


def test_cli_runs():
    p = make(tmp("cli.xlsx"), BASE_IN, BASE_SENS)
    out = subprocess.run([sys.executable, os.path.join(HERE, "check_sens_coverage.py"), p], capture_output=True, text=True)
    assert out.returncode == 0 and "通過" in out.stdout


# ---- 2026-10-30 獨立核對反例（/home/claude/senscov_scripts/cx.py、cx4.py；PR #255 留言 6513910583）----
BETA = ("beta", "Beta", 1.2, 1.5, 2.0)
INS_B = BASE_IN + [BETA]
NO_HE = [x for x in BASE_SENS if x[1] != "he_all"]


def miss(r):
    return {m[0] for m in r["missing"]}


def rebuild_sens(p, top_rows, cases, blank_before=0, keyhdr="鍵"):
    wb = openpyxl.load_workbook(p)
    if "Sens" in wb.sheetnames:
        wb.remove(wb["Sens"])
    s = wb.create_sheet("Sens")
    for t in top_rows:
        s.append([t])
    s.append(["驅動", "設定", "方向", "加權目標價", "相對基準", keyhdr, "說明"])
    for i, (drv, k, n) in enumerate(cases):
        if i == len(cases) - 1:
            for _ in range(blank_before):
                s.append([None])
        s.append([drv, "x", "低", 1, 0, k, n])
    wb.save(p)
    return p


def test_cx1_key_only_in_footnote_row():
    # 頁尾讀法列（沒有目標價數值）不是案例列，提到 beta 不算涵蓋
    p = make(tmp("c1.xlsx"), INS_B, [])
    p = rebuild_sens(p, ["Sens", "註"], BASE_SENS)
    wb = openpyxl.load_workbook(p); wb["Sens"].append(["讀法：龍捲風依絕對值排序；beta 見 Val 頁"]); wb.save(p)
    assert miss(C.check(p)) == {"beta"}
    # 反例原樣：讀法列帶目標價數值（make() 產生），A 欄以「讀法」開頭仍不算案例
    r = C.check(make(tmp("c1o.xlsx"), INS_B, BASE_SENS + [("讀法：龍捲風依絕對值排序；beta 見 Val 頁", "", "")]))
    assert miss(r) == {"beta"}


def test_cx1b_key_only_in_other_case_note():
    r = C.check(make(tmp("c1b.xlsx"), INS_B, BASE_SENS + [("WACC +1pt", "wacc", "beta 維持基準不變")]))
    assert miss(r) == {"beta"}


def test_cx2_key_without_exempt_word():
    assert miss(C.check(make(tmp("c2.xlsx"), INS_B, BASE_SENS, a2="beta：影響小。"))) == {"beta"}


def test_cx2b_negated_exemption():
    r = C.check(make(tmp("c2b.xlsx"), INS_B, BASE_SENS, a2="beta 不得豁免，下版補案例。"))
    assert miss(r) == {"beta"} and "豁免不合格" in r["missing"][0][5]


def test_cx2c_exemption_without_reason():
    r = C.check(make(tmp("c2c.xlsx"), INS_B, BASE_SENS, a2="beta 不列。"))
    assert miss(r) == {"beta"} and "豁免不合格" in r["missing"][0][5]


def test_cx2d_pending_not_exemption():
    assert miss(C.check(make(tmp("c2d.xlsx"), INS_B, BASE_SENS, a2="beta 案例待補，不列入本版龍捲風。"))) == {"beta"}


def test_cx3_single_generation_with_tongshi_text():
    s3 = NO_HE + [("IF_HoldEcon｜GB300 與 k 同時調整：低", "he_b3", "")]
    r = C.check(make(tmp("c3.xlsx"), BASE_IN, s3))
    assert miss(r) == {"he_h", "he_b2", "he_b3", "he_vr"} and "IF_HoldEcon" not in r["group_cov"]


def test_cx3b_pending_note_not_group_case():
    s3b = NO_HE + [("註：IF_HoldEcon 四世代同時案例待補", "", "")]
    assert miss(C.check(make(tmp("c3b.xlsx"), BASE_IN, s3b))) == {"he_h", "he_b2", "he_b3", "he_vr"}


def test_cx3b2_he_all_with_negation_not_group_case():
    s = NO_HE + [("IF_HoldEcon 四世代：待補", "he_all", "")]
    assert miss(C.check(make(tmp("c3b2.xlsx"), BASE_IN, s))) == {"he_h", "he_b2", "he_b3", "he_vr"}


def test_cx3c_three_ranged_generations_three_keys_pass():
    ins = BASE_IN[:3] + HE[:3] + [("he_vr", "IF_HoldEcon｜VR200", 12.2, 12.2, 12.2)]
    s = NO_HE + [("三世代：低", "", "he_h、he_b2、he_b3 同取低")]  # 反例原樣：成員鍵寫在說明欄
    assert not C.check(make(tmp("c3c.xlsx"), ins, s))["missing"]


def test_cx3d_four_ranged_three_keys_fail():
    s = NO_HE + [("三世代：低", "", "he_h、he_b2、he_b3 同取低")]
    assert miss(C.check(make(tmp("c3d.xlsx"), BASE_IN, s))) == {"he_h", "he_b2", "he_b3", "he_vr"}


def test_cx4_header_row_18_found_and_note_not_case():
    p = make(tmp("c4.xlsx"), INS_B, [])
    tops = [f"註記第{i}行：說明文字 he_all 待補" if i == 5 else f"註記第{i}行" for i in range(16)] + [None]
    p = rebuild_sens(p, tops, NO_HE)
    r = C.check(p)
    assert r["keycol"] is True and miss(r) == {"beta", "he_h", "he_b2", "he_b3", "he_vr"}


def test_cx4x_header_beyond_40_rows_fails():
    p = make(tmp("c4x.xlsx"), BASE_IN, [])
    p = rebuild_sens(p, ["註：IF_HoldEcon 四世代同時案例待補（he_all）" if i == 5 else f"註記第{i}行" for i in range(45)], BASE_SENS)
    r = C.check(p)
    assert r["error"] and "表頭" in r["error"] and C.main([p]) == 1


def test_cx4b_case_after_many_blank_rows_counts():
    p = make(tmp("c4b.xlsx"), INS_B, [])
    p = rebuild_sens(p, ["Sens", "註"], BASE_SENS + [("Beta 1.2", "beta", "重算值")], blank_before=11)
    assert not C.check(p)["missing"]


def test_cx4c_keycol_header_variant():
    p = make(tmp("c4c.xlsx"), INS_B, [])
    p = rebuild_sens(p, ["Sens", "註"], BASE_SENS + [("Beta：低", "beta", "")], keyhdr="鍵（In 頁）")
    r = C.check(p)
    assert r["keycol"] is True and not r["missing"]


def test_cx5_loose_not_model_text_not_excluded():
    ins = BASE_IN + [("beta", "Beta", 1.2, 1.5, 2.0, "另一來源的 β 只作參考、不進模型；本列進 WACC")]
    r = C.check(make(tmp("c5.xlsx"), ins, BASE_SENS))
    assert miss(r) == {"beta"} and not r["not_model"]


def test_cx6_one_side_case_accepted():
    assert not C.check(make(tmp("c6.xlsx"), INS_B, BASE_SENS + [("Beta：低", "beta", "")]))["missing"]


def test_cx7_substring_key():
    ins = BASE_IN + [("tax_f", "稅率 F", 0.1, 0.2, 0.3)]
    assert miss(C.check(make(tmp("c7.xlsx"), ins, BASE_SENS))) == {"tax_f"}


# ---- 第 5、6 項：不進模型的引用鏈、IF_HoldEcon 無區間警告、--require-keycol ----
CONS = ("cons_tp", "市場共識目標價", 75, 245, 800, "低／高＝最低／最高（只作對照，不進模型）")


def _cons_row(p):
    ws = openpyxl.load_workbook(p)["In"]
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "cons_tp")


def test_not_model_referenced_by_model_chain_fails():
    p = make(tmp("nm_bad.xlsx"), BASE_IN + [CONS], BASE_SENS)
    r0 = _cons_row(p)
    wb = openpyxl.load_workbook(p)
    pl = wb.create_sheet("PL"); pl["B2"] = f"=In!E{r0}*2"
    wb.save(p)
    r = C.check(p)
    assert miss(r) == {"cons_tp"} and "被計算鏈引用" in r["missing"][0][5] and "PL!B2" in r["missing"][0][5]


def test_not_model_display_only_val_and_recon_ok():
    p = make(tmp("nm_ok.xlsx"), BASE_IN + [CONS], BASE_SENS)
    r0 = _cons_row(p)
    wb = openpyxl.load_workbook(p)
    v = wb.create_sheet("Val"); v["B10"] = f"=In!E{r0}"
    rc = wb.create_sheet("Recon"); rc["C14"] = f"=In!E{r0}"
    wb.save(p)
    r = C.check(p)
    assert not r["missing"] and r["not_model"] == ["cons_tp"]


def test_not_model_val_cell_used_downstream_fails():
    p = make(tmp("nm_val.xlsx"), BASE_IN + [CONS], BASE_SENS)
    r0 = _cons_row(p)
    wb = openpyxl.load_workbook(p)
    v = wb.create_sheet("Val"); v["B10"] = f"=In!E{r0}"; v["B11"] = "=B10*0.5"
    wb.save(p)
    assert miss(C.check(p)) == {"cons_tp"}


def test_holdecon_without_range_warns():
    ins = BASE_IN[:3] + [("tk_hold_h100", "IF_HoldEcon｜Hopper H100", None, 9.7, None)]
    r = C.check(make(tmp("warn.xlsx"), ins, BASE_SENS))
    assert not r["missing"] and len(r["warn"]) == 1 and "tk_hold_h100" in r["warn"][0]


def test_require_keycol():
    # 無鍵欄、案例靠 A 欄文字寫鍵：一般模式通過；--require-keycol 失敗
    sens = [("有效稅率 tax：低", "", ""), ("only_hi：高", "", ""), ("四世代 he_h he_b2 he_b3 he_vr：低", "", "")]
    p = make(tmp("rk.xlsx"), BASE_IN, sens)
    wb = openpyxl.load_workbook(p); wb["Sens"]["F3"] = "代號"; wb.save(p)
    r = C.check(p)
    assert r["keycol"] is False and not r["missing"]
    assert C.main([p]) == 0 and C.main([p, "--require-keycol"]) == 1
    assert C.main([make(tmp("rk2.xlsx"), BASE_IN, BASE_SENS), "--require-keycol"]) == 0

if __name__ == "__main__":
    fs = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for f in fs:
        f()
    print(f"{len(fs)} 項測試通過")
