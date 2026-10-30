# Meta v0.1 成品核對（工作單 B 段第 5 項）：公式錯誤 0、核對列 0（Q、Thesis〔C10、M6〕、Seg、PL〔分部對帳〕、CF、Val）、敏感度未過期、固定列標籤存在。
# 用法：python3 meta/check_v01.py <重算後的 xlsx>；失敗時非零結束。
import sys
import openpyxl

X = sys.argv[1]
wb = openpyxl.load_workbook(X, data_only=True)
fail = []
ERR = ("#REF", "#DIV", "#VALUE", "#NAME", "#N/A", "#NUM", "Err:")
n_err = 0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(ERR):
                n_err += 1
                if n_err <= 10: fail.append(f"公式錯誤 {ws.title}!{c.coordinate} {c.value}")
print(f"公式錯誤：{n_err} 格")


def rows_with(ws, prefix):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith(prefix):
            yield r


def nums(ws, r):
    return [c.value for c in ws[r][1:] if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


n_chk = {}
for sh in ("Q", "Thesis", "Seg", "PL", "CF", "Val"):
    ws = wb[sh]
    n_chk[sh] = 0
    for r in rows_with(ws, "檢查："):
        if "應為 1" in str(ws.cell(r, 1).value): continue
        vs = nums(ws, r)
        n_chk[sh] += 1
        if not vs:
            fail.append(f"{sh}!A{r}「{ws.cell(r, 1).value}」沒有數值")
        if any(abs(v) > 1e-6 for v in vs):
            fail.append(f"{sh}!A{r}「{ws.cell(r, 1).value}」有非零值 {[round(v, 6) for v in vs if abs(v) > 1e-6][:5]}")
print("核對列：" + "、".join(f"{k} {v} 列" for k, v in n_chk.items()) + "（含三情境；Thesis＝C10 與 M6 對帳，PL＝分部對帳）")
for sh, need in (("Thesis", 6), ("PL", 12)):
    if n_chk[sh] < need: fail.append(f"{sh} 核對列少於 {need}（{n_chk[sh]}）")
V = wb["Val"]
need = ["目標價（DCF 與 Comparable 加權）", "上檔空間：加權目標價", "投資建議", "上檔空間級距", "DCF 目標價（情境加權）", "Comparable 目標價（市場主流同業法）",
        "價格帶（DCF 不如預期～順利）", "現價（2026-10-28 收盤）", "市場共識目標價", "讀法 2（自用 AI 價值中性）：目標價／上檔空間／投資建議", "讀法 1（自用 AI 只計成本）：目標價／上檔空間／投資建議"]
labs = {V.cell(r, 1).value: r for r in range(1, V.max_row + 1)}
for n in need:
    if n not in labs: fail.append(f"Val 缺列「{n}」")
cf = wb["CF"]
if "收支關卡：自由現金流扣股利（五年合計）" not in {cf.cell(r, 1).value for r in range(1, cf.max_row + 1)}:
    fail.append("CF 缺「收支關卡：自由現金流扣股利（五年合計）」")
se = wb["Sens"]
for r in range(1, se.max_row + 1):
    if str(se.cell(r, 1).value or "").startswith("敏感度重算時的基準值"):
        d = se.cell(r, 5).value
        if not isinstance(d, (int, float)) or abs(d) > 1e-6:
            fail.append(f"Sens 敏感度已過期或缺：基準差 {d}")
tp = V.cell(labs["目標價（DCF 與 Comparable 加權）"], 2).value
up = V.cell(labs["上檔空間：加權目標價"], 2).value
rec = V.cell(labs["投資建議"], 2).value
r1 = labs["讀法 1（自用 AI 只計成本）：目標價／上檔空間／投資建議"]
print(f"加權目標價 {tp:.2f}；上檔空間 {up:+.1%}；投資建議 {rec}（讀法 1：{V.cell(r1, 2).value:.2f}，{V.cell(r1, 4).value}）")
if fail:
    print("失敗：")
    for f in fail: print("  -", f)
    sys.exit(1)
print("全部通過")
