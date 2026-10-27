# 勝宏科技（300476.SZ／02476.HK）收支與估值模型 v0.1.1（AI 半導體 Project：載板與 PCB 次產業；AI 加速卡與伺服器高層數 PCB、HDI）
# 用法：python3 victory_giant/build_victory_giant.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億人民幣（RMB B）；每 GW PCB 金額以美元輸入、以匯率 fx 換算；EPS：人民幣。
# 口徑：歸屬母公司淨利（中國企業會計準則）。樣板：emc/（台光電：每 GW 高速 CCL 金額 × 份額，產能上限）。
# 勝宏在 CCL 的下游（PCB）：每 GW PCB 金額在本模型 Units 頁另立（與台光電的每 GW CCL 金額口徑分開），Recon 頁兩者對照。
# FY2030–FY2032 為收斂法（上游只到 2029）；目標價＝DCF 與 Comparable 加權（Val 第 0 節）。
import sys, os, glob
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "siliconomics"))
import reader as SIL   # Siliconomics 快照（行業變數單一來源）
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import cycle_position as CP
import valuation_rule as VR

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
NET_F = only("markets/network/dist/*.xlsx"); TSMC_F = only("tsmc/dist/*.xlsx"); BR_F = only("markets/chip_bridge/dist/*.xlsx")
EMC_F = only("emc/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
ITC = None
for r in openpyxl.load_workbook(NET_F, data_only=True)["Mkt"].iter_rows(min_row=4):
    if r[0].value == "AI IT 資本支出（加速器 × 每顆 IT capex）":
        ITC = ([r[2 + j].value for j in range(4)], f"{rel(NET_F)}｜Mkt!C{r[0].row}:F{r[0].row}")
assert ITC and all(v is not None for v in ITC[0]), ITC
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith("um"):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
assert all(f"um{y}_{s}" in TS for y in (28, 29) for s in SC), TS.keys()
BR = {}
for r in openpyxl.load_workbook(BR_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k in ("racks", "capit"):
        BR[k] = (r[4].value, r[7].value, f"{rel(BR_F)}｜In!E{r[0].row}（GB300）、H{r[0].row}（VR200）")
assert len(BR) == 2, BR.keys()
EMC = {}
for r in openpyxl.load_workbook(EMC_F, data_only=True)["Units"].iter_rows(min_row=4):
    lab = r[0].value
    for key, want in (("ccl", "每 GW 高速 CCL 金額（合計，不含價格指數）"), ("emcpg", "每 GW 台光電 AI 相關營收")):
        if lab == want:
            EMC[key] = ([r[2 + j].value for j in range(4)], f"{rel(EMC_F)}｜Units!C{r[0].row}:F{r[0].row}")
assert len(EMC) == 2 and all(v is not None for v in EMC["ccl"][0]), EMC

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億人民幣（RMB B），每 GW PCB 金額為百萬美元；比例：%。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間。SemiAnalysis 未使用。公司原文＝勝宏科技定期報告（深交所、巨潮）與業績交流會（經媒體轉述）；公司的訂單、產品組合與產能說法為 Interested-party；券商估計皆為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
H1R = "https://disc.static.szse.cn/disc/disk03/finalpage/2026-08-28/296a3692-b153-48c1-9201-2bbb41886183.PDF"
WS1H = "https://wallstreetcn.com/articles/3780462"
IFH = "https://finance.ifeng.com/c/8wJ1y4oRRjn"
Q1A = "https://www.itiger.com/news/1110866401"
CALL = "https://news.10jqka.com.cn/20260828/c679385007.shtml"
FY25 = "https://stock.stockstar.com/RB2026031400003895.shtml"
WS25 = "https://wallstreetcn.com/articles/3767370"
JM = "https://m.jiemian.com/article/14114673.html"
JPM = "https://finance.sina.com.cn/stock/hkstock/marketalerts/2026-06-11/doc-iniazihq5543947.shtml"
GS = "https://news.qq.com/rain/a/20260903A09PBG00"
GJ = "https://pdf.dfcfw.com/pdf/H3_AP202605241822839616_1.pdf"
THS = "https://basic.10jqka.com.cn/300476/worth.html"
SA = "https://stockanalysis.com/quote/she/300476/"
CAP = "https://caifuhao.eastmoney.com/news/20260105193303948073610"
RUM = "https://finance.sina.com.cn/wm/2026-07-21/doc-iniiqvau9153021.shtml"
WK = "https://waikanstore.com/victory-giant-300476-company-research-2026-10/"
FXS = "https://cj.sina.com.cn/articles/view/1663612603/6328b6bb0200248kg"

sec("A. AI 部署規模（讀 markets/network、markets/chip_bridge、tsmc 成品；情境 A 基準，B／C 依顆數係數縮放）")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"itc_{y}", f"AI IT 資本支出 20{y}", "US$B", None, ITC[0][j], None, "Derived", ITC[1],
        "加速器顆數 × 每顆 IT capex（含 ASIC）；本模型以此除以每 GW IT 資本支出得到部署 GW（與台光電、台達電模型相同）" if y == 26 else "同上")
inp("itpg_gb", "每 GW IT 資本支出（GB300）", "US$B", None, BR["capit"][0], None, "Derived", BR["capit"][2], "Tokenomics IF_CapexIT（經橋接表；唯讀）")
inp("itpg_vr", "每 GW IT 資本支出（VR200；2027 起）", "US$B", None, BR["capit"][1], None, "Derived", BR["capit"][2], "同上；Rubin Ultra 與 ASIC 機架未另訂（Assumed：同 VR200）")
inp("vr26", "2026 年部署中 Rubin 世代占比（混合每 GW IT 資本支出與每 GW PCB 金額用）", "%", *SIL.vals("vr26"), "Interested-party", "TrendForce 2026 組合 Rubin 22%（經 markets/memory；與台光電、台達電模型相同）", "VR200 自 2026 Q3 起出貨（國金證券 2026-05，Interested-party）")
inp("racks_gb", "每 GW 機架數（GB300）", "架", None, BR["racks"][0], None, "Derived", BR["racks"][2], "Recon 用：每架 PCB 金額換算")
inp("racks_vr", "每 GW 機架數（VR200）", "架", None, BR["racks"][1], None, "Derived", BR["racks"][2], "同上")
for s in SC:
    inp(f"um28_{s}", f"情境 {s}：2028 顆數 ÷ 基準", "x", None, TS[f"um28_{s}"][0], None, "Assumed", TS[f"um28_{s}"][1], "A＝順利（基準）；B＝平平（AI 資本支出消化）；C＝不如預期（明顯下修）；與台積電、台光電模型相同")
    inp(f"um29_{s}", f"情境 {s}：2029 顆數 ÷ 基準", "x", None, TS[f"um29_{s}"][0], None, "Assumed", TS[f"um29_{s}"][1], "同上")
inp("fx", "匯率（人民幣 ÷ 美元）", "RMB/US$", 6.50, 6.72, 7.00, "Verified", "中國經濟網 2026-10-16：中間價 6.7220（網址見 Sources 頁）", "10-09 中間價 6.7330（中國經濟網）；10-23 中間價：未找到")

sec("B. 每 GW AI 高速 PCB 金額與勝宏份額（本專案權威：每 GW PCB 金額、勝宏份額，規則第 13 條；與台光電的每 GW CCL 金額口徑分開）")
TOPDOWN = ("由上而下（Assumed）：三項合計 2026 年約 US$450M／GW；量級檢查：Morgan Stanley 每架 PCB GB300 US$35.1K、VR200 US$116.7K（國金證券 2026-05 轉引，Interested-party）"
           "× 每 GW 架數＝237／473 百萬美元（只含 NVL72 機架內，不含 scale-out 交換器、光模組、ASIC 伺服器）；Prismark AI／HPC PCB 2024 年 US$6B（摩根大通 2026-06-11 轉引）")
inp("cmp_26", *SIL.row9("pcb_cmp_26")[1:])
inp("sw_26", *SIL.row9("pcb_sw_26")[1:])
inp("mid_26", *SIL.row9("pcb_mid_26")[1:])
IDX = (("cmpx", "加速器卡／運算板", ((1.25, 1.50, 1.75), (1.35, 1.65, 2.00), (1.40, 1.80, 2.20)), "Analogy",
        "Morgan Stanley 每架 PCB（國金證券轉引）：VR200 ÷ GB300＝3.3 倍；× 每 GW 架數（4,052 ÷ 6,764）＝每 GW 2.0 倍；2026 年已含 22% Rubin → 2027 ÷ 2026 約 1.64 倍（含交換匣，本列取 1.50）",
        "Bernstein：AI 伺服器 PCB 含量較前一代約倍增（國金證券轉引，Interested-party）"),
       ("swx", "交換器板", ((1.15, 1.35, 1.60), (1.25, 1.50, 1.85), (1.30, 1.65, 2.10)), "Assumed",
        "交換匣 24 層 → 32 層、1.6T 2027 年放量（國金證券、公司）", "網通模型：交換器市場 ÷ AI IT 資本支出 2026–27 持平；本列的上升來自層數與材料等級"),
       ("midx", "中板／背板／其他", ((1.40, 2.00, 2.60), (1.80, 3.00, 4.00), (2.00, 3.50, 5.00)), "Assumed",
        "VR200 中板與 CPX（2026 Q3 起）、Kyber 78 層正交背板（2027 起）；國金證券：Kyber 每架 PCB「十萬美元級」以上", "台光電模型 UBB／中板指數 2027 1.80、2028 2.30、2029 2.60（CCL 口徑）"))
for key, lab, vals, tag, s1, s2 in IDX:
    for y in (27, 28, 29):
        inp(f"{key}_{y}", *SIL.row9(f"pcb_{key}_{y}")[1:])
for y, v in ((27, (0.70, 0.90, 1.00)), (28, (0.60, 0.80, 1.00)), (29, (0.55, 0.75, 0.95)), ):
    inp(f"rel_{y}", f"勝宏份額 20{y} ÷ 2026（相對值）", "x", *v, "Assumed",
        "摩根大通：勝宏 AI／HPC PCB 份額 2024 年 1.7% → 1H25 13.8%（第一）、高階 HDI 44.6%（Prismark，Interested-party）；市場傳聞 Rubin 份額 55% → 30%（2026-07，公司否認，未證實）" if y == 27 else "同上",
        "對手：沪电股份、深南电路、生益电子、鹏鼎控股、景旺电子、臺灣與美國廠（TTM 等）擴產；勝宏 1H26 營收 +28.8%，低於沪电 +61%、深南 +46%（外刊庫 2026-10）" if y == 27 else "—")
for y, v in ((27, (0.95, 1.00, 1.03)), (28, (0.90, 0.97, 1.02)), (29, (0.85, 0.94, 1.00))):
    inp(f"pp_{y}", f"PCB 價格指數 20{y}（2026＝1；同規格降價或漲價，不含世代升級）", "x", *v, "Assumed",
        "市場傳聞 NVIDIA 要求 PCB 廠統一降價 10%（2026-06，NVIDIA 否認；新浪 2026-07-21）" if y == 27 else "同上",
        "公司：新一代 PCB 單價提升數倍（世代升級，已在每 GW 金額指數內）；高多層與高階 HDI 排單到 2027（公司 2026-08-28）" if y == 27 else "新產能到位後價格回落（Assumed）")

sec("C. 產能（全產全銷：營收上限＝FY2026 營收 × 產值產能指數 × 價格指數）")
CAPS = "公司：2030 年千億（1,000 億元）產值目標；在產＋在建＋規劃產能約等於該目標（界面 2026-03）；2026 年固定資產投資不超過 180 億元（公司 2026-08-28）"
CAPD = "東方財富財富號（2026-01-05，產能整理，Interested-party）：惠州廠房 10、11 各約 100 萬㎡、產值各逾 150 億（2026 年底投產）、廠房 12 約 120 億（2027）、越南一期 30–35 億（2026 年中）、泰國 A2 36 億（2026）、泰國二期 45 萬㎡（2027）"
for y, v in ((27, (1.45, 1.70, 1.90)), (28, (1.95, 2.40, 2.80)), (29, (2.40, 3.00, 3.40))):
    inp(f"cx_{y}", f"平均產值產能 20{y} ÷ FY2026 營收（產能指數；以 FY2026 產品價格計）", "x", *v, "Assumed",
        CAPS if y == 27 else ("同上；2029 年 × FY2026 營收約 840 億，向 2030 年千億收斂" if y == 29 else "同上"),
        CAPD if y == 27 else ("新廠爬坡：4Q25、1Q26 低於預期主因新產能爬坡慢、原料上漲（財信證券，經界面 2026-03）；低值反映延遲" if y == 28 else "—"))

sec("D. 營收結構錨點（公司揭露比例 × 營收）")
inp("ai_25", "2025 年 AI 相關營收（AI 算力、資料中心、HPC）÷ 營收", "%", 0.40, 0.45, 0.50, "Interested-party", "摩根大通 2026-06-11：AI／HPC 占營收 2024 年 6.6% → 2025 年 43.2%（招股書口徑）", "公司：2025 前三季 AI 相關超過 40%、全年更高（界面 2026-03）；業績交流會：去年 AI 產品約占一半（同花順 2026-08-28）")
inp("g_nai26", "2026 年非 AI 營收年增（公司：非 AI 金額大致持平）", "%", -0.05, 0.0, 0.05, "Interested-party", CALL + "：預期非 AI 營收絕對額大致持平，成長來自 AI", "FY2026 AI 相關營收＝錨點營收 − 非 AI（Derived）")
for y, v in ((27, (-0.05, 0.0, 0.05)), (28, (-0.05, 0.0, 0.05)), (29, (-0.05, 0.0, 0.05))):
    inp(f"g_nai{y}", f"非 AI 營收年增 20{y}（汽車、手機、PC、消費、工控）", "%", *v, "Assumed",
        "公司：非 AI 金額持平（2026-08-28）" if y == 27 else "同上", "產能優先給 AI；新產能另含長沙、益陽單雙面板（界面 2026-03）" if y == 27 else "—")

sec("E. 毛利率")
inp("d_ai", "AI 相關（高多層、高階 HDI）毛利率高於其他營收的幅度", "pt", 0.10, 0.15, 0.20, "Assumed", "部門別毛利率未揭露：未找到；摩根大通：HDI 毛利率 2025 年 43.5%、多層板 24.4%（Interested-party）", "公司：HDI 毛利率 2024 年 22.5% → 2025 前三季 45%（界面 2026-03）")
for y, v in ((27, (-0.01, 0.01, 0.03)), (28, (-0.01, 0.02, 0.04)), (29, (-0.01, 0.03, 0.05))):
    inp(f"gmu_{y}", f"AI 相關毛利率額外改善 20{y}（新廠爬坡完成、新世代產品組合；不含價格）", "pt", *v, "Assumed",
        "公司：2Q26 毛利率承壓來自上半年人員增加逾 30%、新廠設備折舊（2026-08-28）；毛利率 2Q25 38.8% → 4Q25 33.5% → 2Q26 約 32.0%" if y == 27 else "同上",
        "摩根大通：毛利率 2026 38.3%、2027 41.3%、2028 42.5%（Interested-party，只作 Recon）" if y == 27 else "—")

sec("F. FY2026 錨點（1H26 實績＋3Q、4Q26 假設；3Q26 季報預計 2026-10-29 公布，截至 2026-10-25 尚未公布）")
inp("q1_rev", "1Q26 營收（實績）", "RMB B", None, 5.519, None, "Verified", Q1A + "：55.19 億元、年增 27.99%", "1Q26 歸母淨利 12.88 億元（同）；1H26 營收 116.29 億（公司半年報）→ 2Q26 61.10 億")
inp("q3g", "3Q26 營收季增", "%", 0.10, 0.20, 0.35, "Assumed", "公司：2Q 存貨增加約 20 億元，為下半年新產品大規模上量備貨；VR200 2026 Q3 起出貨（國金證券）", "共識 FY2026 營收 316 億（同花順 14 家，2026-10-24）隱含下半年 200 億、較上半年 +72%（只作 Recon）")
inp("q4g", "4Q26 營收季增", "%", 0.10, 0.25, 0.40, "Assumed", "惠州廠房 10、11 2026 年底投產；越南一期、泰國 A2 爬坡（東方財富財富號 2026-01）", "摩根大通：4Q26 淨利年增 217%（Interested-party）")
inp("q_gm", "3Q、4Q26 毛利率", "%", 0.30, 0.33, 0.36, "Assumed", "1H26 33.2%（營業成本 77.66 億、營收 116.29 億；公司半年報）；1Q26 34.5%、2Q26 約 32.0%（Derived）", "新產品上量與新廠爬坡相抵（公司 2026-08-28）")
inp("q_opx", "3Q、4Q26 營業費用（銷售＋管理＋研發）÷ 營收", "%", 0.060, 0.065, 0.070, "Assumed", "1H26：銷售 1.05＋管理 2.92＋研發投入 3.59 億＝7.56 億、6.5%（公司半年報；研發投入是否全數費用化未載明）", "—")
inp("q_nop", "3Q、4Q26 營業外與財務淨額（每季；財務費用、公允價值變動、稅金及附加等）", "RMB B", -0.12, -0.06, 0.06, "Assumed", "1H26 淨額 +1.86 億（稅前 − 毛利 + 營業費用；含公允價值變動 +5.95 億、財務費用 −2.97 億）", "H 股募資約 200 億港元、交易性金融資產 153 億元（鳳凰網 2026-08）；公允價值變動不外推")
inp("q_nir", "歸屬母公司淨利 ÷ 稅前（3Q、4Q26）", "%", 0.84, None, 0.88, "Derived", "1H26：歸母 28.57 億 ÷（歸母 28.57＋所得稅 4.37）＝0.867（公司半年報；少數股東損益未載明，視為 0）", "高新技術企業稅率 15%（Assumed）")
inp("sh_1h", "1H26 加權股數（十億股）", "十億股", None, "=2.8568/3.14", None, "Derived", "1H26 歸母淨利 28.57 億 ÷ 基本 EPS 3.14（公司半年報）", "H 股 2026-04-21 上市，1.10 億股（含超額配售）")
inp("sh_h2", "下半年股數（A 股 8.7234 億＋H 股 1.10 億；十億股）", "十億股", None, "=0.87234+0.110", None, "Derived", "A 股：2025 年度分派基數 872,339,870 股（公司半年報）；H 股 1.10 億股（華爾街見聞 2026-08-27）", "stockanalysis：總股數 982.57 百萬股")

sec("G. 營業費用、業外、稅與股數（FY2027 起）")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.40, 0.60, 1.00, "Assumed", "1H26 營業費用（不含財務費用）年增 7%（7.56 ÷ 7.02 億）、營收 +28.8%（公司半年報）", "上半年人員增加逾 30%（公司）；新廠開出期間費用可能較快")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.01, 0.0, 0.01, "Assumed", "無", "Sens 用 ±1 個百分點")
inp("nop_f", "營業外與財務淨額（每年）", "RMB B", -0.5, -0.2, 0.2, "Assumed", "借款：短期 32.4 億、長期 59.9 億元（1H26，鳳凰網）；交易性金融資產 153 億", "公允價值變動、匯兌不外推")
inp("nir_f", "歸屬母公司淨利 ÷ 稅前（FY2027 起）", "%", 0.84, 0.867, 0.88, "Derived", "同 q_nir（1H26 0.867）", "—")
inp("sh_g", "股數年變動", "%", 0.0, 0.0, 0.01, "Assumed", "限制性股票不超過 486.24 萬股（財聯社 2026-06）", "—")
inp("mf", "AI 部署係數（Sens 用；FY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "網通模型 2027 年 AI IT 資本支出約 2 倍", "±20%")

sec("H. 共識、券商、研調與上游 CCL（只作 Recon 對帳，不作校準）")
for y, rv, ep in ((26, 31.604, 8.61), (27, 51.477, 14.77), (28, 74.435, 22.13)):
    inp(f"cons_rev{y}", f"共識 FY20{y} 營收（同花順 14 家平均）", "RMB B", None, rv, None, "Interested-party", THS + "（2026-10-24 擷取；最新報告 2026-09-07）", "—")
    inp(f"cons_eps{y}", f"共識 FY20{y} EPS（同花順 14 家平均）", "RMB", None, ep, None, "Interested-party", THS, {26: "區間 7.18–10.00", 27: "區間 10.42–17.45", 28: "區間 14.70–26.17"}[y])
for y, rv, ep in ((26, 35.5, 9.20), (27, 60.0, 17.55), (28, 84.2, 26.06)):
    inp(f"jpm_rev{y}", f"摩根大通 FY20{y} 營收（2026-06-11）", "RMB B", None, rv, None, "Interested-party", JPM, "—")
    inp(f"jpm_eps{y}", f"摩根大通 FY20{y} EPS（2026-06-11）", "RMB", None, ep, None, "Interested-party", JPM, "目標價 600 港元＝24 倍 2026–27 平均 P/E" if y == 26 else "—")
inp("gs_eps27", "高盛隱含 FY2027 EPS（目標價 550 ÷ 26.3 倍 2027 P/E）", "RMB", None, "=550/26.3", None, "Derived", GS + "（2026-09-03，Interested-party）", "—")
inp("jpm_sh", "摩根大通：勝宏 AI／HPC PCB 份額 1H25（Prismark）", "%", None, 0.138, None, "Interested-party", JPM, "2024 年 1.7%；高階 HDI 44.6%")
inp("ms_ccl_gb", "Morgan Stanley：每架 CCL 金額 GB300（台光電模型所引）", "US$K", None, 6.695, None, "Interested-party", "https://www.itiger.com/news/1142209073（Deep News 2026-09-23 轉述；與台光電模型 In 頁相同）", "VR200 18.75K；用於對照每架 PCB ÷ 每架 CCL")
inp("ms_rk_gb", "Morgan Stanley：每架 PCB 金額 GB300 NVL72", "US$K", None, 35.1, None, "Interested-party", GJ + "（國金證券轉引 Morgan Stanley 供應鏈調查）", "整櫃 BOM 3.995 百萬美元")
inp("ms_rk_vr", "Morgan Stanley：每架 PCB 金額 VR200 NVL72", "US$K", None, 116.7, None, "Interested-party", GJ, "+233%；整櫃 BOM 7.803 百萬美元")
for j, y in enumerate((26, 27, 28, 29)):
    inp(f"emc_ccl{y}", f"台光電模型：每 GW 高速 CCL 金額 20{y}（合計）", "US$M/GW", None, EMC["ccl"][0][j], None, "Derived", EMC["ccl"][1], "唯讀台光電成品（CCL 口徑；與本模型每 GW PCB 金額分開，只在 Recon 對照）" if y == 26 else "同上")
for j, y in enumerate((26, 27, 28)):
    inp(f"emc_pg{y}", f"台光電模型：每 GW 台光電 AI 相關營收 20{y}", "US$M/GW", None, EMC["emcpg"][0][j], None, "Derived", EMC["emcpg"][1], "唯讀台光電成品")

sec("I. 價格與估值")
inp("px", "股價（2026-10-23 收盤，深交所 A 股）", "RMB", None, 183.48, None, "Interested-party", SA + "：10-23 收盤 183.48（+3.08%）、前收 177.99", "搜狐 2026-10-23：當日 +3.08%、總市值 1,803.21 億元")
inp("pe_ttm", "近 12 月 P/E（stockanalysis，10-23）", "x", None, 35.87, None, "Interested-party", SA, "近 12 月 EPS 5.11")
inp("fpe", "前瞻 P/E（stockanalysis，10-23）", "x", None, 14.84, None, "Interested-party", SA, "—")
inp("cons_feps", "共識前瞻 12 月 EPS（現價 ÷ 前瞻 P/E）", "RMB", None, "=px/fpe", None, "Derived", "stockanalysis 前瞻 P/E 14.84（2026-10-23）", "共識前瞻 EPS 的直接數字：未找到")
inp("cons_tp", "市場共識目標價（賣方平均）", "RMB", None, 388.30, None, "Interested-party", SA + "：8 位分析師平均 388.30（2026-10-23 擷取）", "stockanalysis 9 月 27 日：7 位平均 416.68；高盛 550（09-03）、摩根大通 600 港元（06-11）")
REV10 = [1.818, 2.442, 3.304, 3.885, None, 7.432, 7.885, 7.931, 10.731, 19.292]
NI10 = [0.232, 0.282, 0.380, 0.463, None, 0.670, 0.791, 0.671, 1.154, 4.312]
inp("nim_avg", "十年平均淨利率（2016–2025，歸屬母公司淨利 ÷ 營收；2020 年缺，9 年）", "%", None,
    "=AVERAGE(" + ",".join(f"{n}/{v}" for n, v in zip(NI10, REV10) if v) + ")", None, "Derived",
    "2016 民生證券 2017-10-27、2017 東吴證券 2018-10-28、2018 民生 2020-02-28、2019 民生 2020-08-19；2021–2025 stockanalysis 年報數字",
    "2020 年營收與淨利：未找到（付費）；單一來源")

sec("J. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%（Val 頁第 0 節列示）")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%，各市場同一值）", "DCF 折現；情境價值推到 12 個月後")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "DCF 永續成長與 FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均，FY2023–25 加權平均 ROE）", "%", None, "=AVERAGE(0.0935,0.1395,0.3556)", None, "Derived", "公司：FY2025 35.56%、FY2024 13.95%（華爾街見聞 2026-03 年報）；FY2023 9.35%（2024 業績快報，巨潮）", "1H26 加權 ROE 11.80%（半年）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2029 淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（情境 B FY2029 與十年平均的中點）", "%", None, 0, None, "Derived", "PL 頁情境 B FY2029、In 頁 nim_avg", "—")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（＝十年平均淨利率）", "%", None, "=nim_avg", None, "Derived", "In 頁 nim_avg", "—")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    if key == "q_nir": base = "=2.8568/(2.8568+0.4372)"
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    if key in TP_KEYS:
        for j in range(1, 10): IN.cell(r, j).fill = KEY
        IN.cell(r, 7, tag + "（Andy 可改）" if tag == "Assumed" else tag)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for col, w in zip("ABCDEFGHI", [12, 60, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點 ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億人民幣；EPS 為人民幣；股數為十億股）。藍字＝輸入；黑字＝公式。實績取自勝宏科技定期報告（深交所）與年報摘要（經轉述）"
Q["A2"] = ("FY2026 錨點＝1H26 實績＋3Q、4Q 假設（營收季增、毛利率、營業費用率、營業外、淨利 ÷ 稅前）。1H26 營業外淨額＝稅前 − 營業利益（含公允價值變動 +5.95 億，不外推）。"
           "FY2025 營業費用＝毛利 − 營業利益（stockanalysis，含財務費用；1H26 起不含財務費用，口徑不同）。3Q26 季報預計 2026-10-29 公布（截至 2026-10-25 尚未公布），公布後更新")
hdr(Q, 3, ["項目", "FY2025A", "1H26（實績）", "3Q26（模型）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1", "來源2／說明"])
S1 = "FY2025：證券之星 2026-03-14 年報簡析（營收 192.92 億、歸母 43.12 億、毛利率 35.22%、EPS 5.01）；stockanalysis（營業利益 51.48 億）；1H26：公司半年報（深交所 2026-08-28）"
S2 = "1H26：華爾街見聞 2026-08-27、鳳凰網（同值）"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2=S2, ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"F{r}"] = ann or f"=SUM(C{r}:E{r})"; Q[f"F{r}"].number_format = fmt
    Q.cell(r, 7, tag); Q.cell(r, 8, s1); Q.cell(r, 9, s2)
qrow(4, "營收", {"B": 19.292, "C": 11.629, "D": "=(C4-q1_rev)*(1+q3g)", "E": "=D4*(1+q4g)"}, tag="Verified／Assumed（3Q、4Q 季增）", s2="1H26 116.29 億（年增 28.77%）；2Q26＝116.29 − 1Q26 55.19＝61.10 億")
qrow(5, "毛利", {"B": "=B4*0.3522", "C": "=11.629-7.766", "D": "=D4*q_gm", "E": "=E4*q_gm"}, tag="Verified／Assumed", s2="1H26 營業成本 77.66 億（公司半年報）；FY2025 毛利率 35.22%（證券之星）")
qrow(6, "毛利率", {c: f"={c}5/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="1H26 33.2%；1Q26 34.5%（營業成本 36.18 億）")
qrow(7, "營業費用", {"B": "=B5-B8", "C": "=0.1047+0.2917+0.3592", "D": "=D4*q_opx", "E": "=E4*q_opx"}, tag="Derived／Verified／Assumed", s2="1H26：銷售 1.047＋管理 2.917＋研發投入 3.592 億（不含財務費用）；FY2025＝毛利 − 營業利益（含財務費用）")
qrow(8, "營業利益（毛利 − 營業費用）", {"B": 5.148, "C": "=C5-C7", "D": "=D5-D7", "E": "=E5-E7"}, tag="Interested-party（stockanalysis）／Derived", s2="FY2025 stockanalysis 營業利益 51.48 億")
qrow(9, "營業利益率", {c: f"={c}8/{c}4" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(10, "營業外與財務淨額", {"C": "=C11-C8", "D": "=q_nop", "E": "=q_nop"}, tag="Derived／Assumed", s1="1H26＝稅前 − 營業利益（含公允價值變動 +5.95 億、財務費用 −2.97 億）", s2="FY2025：未找到")
qrow(11, "稅前利益", {"C": "=2.8568+0.4372", "D": "=D8+D10", "E": "=E8+E10"}, tag="Derived", s1="1H26＝歸母淨利 28.57＋所得稅 4.37 億（少數股東損益未載明）", s2="FY2025：未找到")
Q["F10"] = "=SUM(C10:E10)"; Q["F11"] = "=SUM(C11:E11)"
qrow(12, "歸屬母公司淨利", {"B": 4.312, "C": 2.8568, "D": "=D11*q_nir", "E": "=E11*q_nir"}, tag="Verified／Derived（3Q、4Q）", s2="FY2025 43.12 億（公司年報）；1H26 28.57 億（年增 33.30%；扣非 24.18 億、+12.50%）")
qrow(13, "淨利 ÷ 營業利益", {c: f"={c}12/{c}8" for c in "BCDEF"}, fmt="0.0%", ann="", tag="Derived", s1="—", s2="—")
qrow(14, "股數（十億股；加權）", {"B": "=4.312/5.01", "C": "=sh_1h", "D": "=sh_h2", "E": "=sh_h2"}, ann="=(sh_1h+sh_h2)/2", tag="Derived", s2="淨利 ÷ EPS；FY2026＝上半年與下半年平均（H 股 4 月下旬上市，Assumed）")
qrow(15, "EPS", {c: f"={c}12/{c}14" for c in "BCDE"}, fmt="0.00", ann="=F12/F14", tag="Derived", s1="淨利 ÷ 股數", s2="—")
qrow(16, "  公司揭露 EPS", {"B": 5.01, "C": 3.14}, fmt="0.00", ann="", tag="Verified", s2="FY2025 基本 5.01；1H26 基本 3.14（公司）")
qrow(17, "AI 相關營收（FY2025＝營收 × AI 占比；FY2026＝營收 − 非 AI）", {"B": "=B4*ai_25"}, ann="=F4-F18", tag="Interested-party／Derived", s1="AI 占比（In 頁 D 節）", s2="各季 AI 營收金額未揭露：未找到")
qrow(18, "非 AI 營收（FY2026＝FY2025 ×（1＋g_nai26））", {"B": "=B4-B17"}, ann="=B18*(1+g_nai26)", tag="Derived／Interested-party", s1="—", s2="公司：非 AI 金額大致持平（2026-08-28）")
Q.cell(23, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
Q.cell(24, 1, "FY2026 營收年增"); Q["B24"] = "=F4/B4-1"
Q.cell(25, 1, "FY2026 AI 相關 ÷ 營收"); Q["B25"] = "=F17/F4"
Q.cell(26, 1, "FY2026 營業費用 ÷ 營收（不含財務費用）"); Q["B26"] = "=F7/F4"
Q.cell(27, 1, "FY2026 營收 ÷ 共識（同花順）− 1"); Q["B27"] = "=F4/cons_rev26-1"
Q.cell(28, 1, "FY2026 EPS ÷ 共識（同花順）− 1"); Q["B28"] = "=F15/cons_eps26-1"
Q.cell(29, 1, "下半年營收 ÷ 上半年 − 1"); Q["B29"] = "=(D4+E4)/C4-1"
for rr in range(24, 30): Q[f"B{rr}"].number_format = "0.0%"
for nm, rr in (("rev", 4), ("gp", 5), ("opx", 7), ("oi", 8), ("ni", 12), ("sh", 14), ("eps", 15), ("ai", 17), ("nai", 18)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$F${rr}")
Q.column_dimensions["A"].width = 60
for c in "BCDEF": Q.column_dimensions[c].width = 15
Q.column_dimensions["G"].width = 26; Q.column_dimensions["H"].width = 90; Q.column_dimensions["I"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎 ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I", 2032: "J"}
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
BU = (2026, 2027, 2028, 2029)
umx = lambda y, s: {2026: "1", 2027: "1", 2028: f"um28_{s}", 2029: f"um29_{s}"}[y]
DEFP = dict(scn="A", cmpx="cmpx", swx="swx", midx="midx", rel="rel", pp="pp", cx="cx", mf="1", gn="g_nai", d="d_ai", gmu="gmu", opxa="0", phi="phi")
YH = ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, YH); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 11, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    c_ = lambda y: FC[y]
    sec_("A. AI 部署與每 GW AI 高速 PCB 金額（需求面；D 欄＝FY2026 校準年）")
    row("itc", "AI IT 資本支出 × 情境顆數係數", "US$B", {c_(y): f"=itc_{y % 100}*{umx(y, s)}" + ("" if y == 2026 else f"*{P['mf']}") for y in BU})
    row("itpg", "每 GW IT 資本支出", "US$B", {"D": "=itpg_gb*(1-vr26)+itpg_vr*vr26", **{c_(y): "=itpg_vr" for y in BU[1:]}})
    row("gw", "部署 GW（IT 電力）", "GW", {c_(y): f"={c_(y)}{R['itc']}/{c_(y)}{R['itpg']}" for y in BU}, fmt="0.0")
    for k, lab, base, ix in (("vcmp", "每 GW PCB：加速器卡／運算板", "cmp_26", P["cmpx"]), ("vsw", "每 GW PCB：交換器板", "sw_26", P["swx"]), ("vmid", "每 GW PCB：中板／背板／其他", "mid_26", P["midx"])):
        row(k, lab, "US$M/GW", {"D": f"={base}", **{c_(y): f"={base}*{ix}_{y % 100}" for y in BU[1:]}}, fmt="#,##0")
    row("pgw", "每 GW AI 高速 PCB 金額（三項合計）", "US$M/GW", {c_(y): f"={c_(y)}{R['vcmp']}+{c_(y)}{R['vsw']}+{c_(y)}{R['vmid']}" for y in BU}, fmt="#,##0")
    row("mku", "AI 高速 PCB 市場（美元；不含價格指數）", "US$B", {c_(y): f"={c_(y)}{R['gw']}*{c_(y)}{R['pgw']}/1000" for y in BU})
    row("mkn", "AI 高速 PCB 市場（人民幣）", "RMB B", {c_(y): f"={c_(y)}{R['mku']}*fx" for y in BU}, fmt="#,##0.0")
    row("pp", "PCB 價格指數（FY2026＝1）", "x", {"D": "=1", **{c_(y): f"={P['pp']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh0", "勝宏有效份額（FY2026 校準）＝ AI 相關營收 ÷ 市場", "%", {"D": f"=a26_ai/D{R['mkn']}"}, fmt="0.0%",
        note="Derived：分子為 FY2026 錨點營收 − 非 AI（公司：非 AI 持平）；分母為本模型的 AI 高速 PCB 市場（含 ASIC）")
    row("rel", "份額相對值（÷ FY2026）", "x", {"D": "=1", **{c_(y): f"={P['rel']}_{y % 100}" for y in BU[1:]}}, fmt="0.00")
    row("sh", "勝宏有效份額", "%", {c_(y): f"=$D${R['sh0']}*{c_(y)}{R['rel']}" for y in BU}, fmt="0.0%")
    row("aid", "AI 相關營收（需求面）＝ 市場 × 份額 × 價格指數", "RMB B", {c_(y): f"={c_(y)}{R['mkn']}*{c_(y)}{R['sh']}*{c_(y)}{R['pp']}" for y in BU}, fmt="#,##0.0")
    sec_("B. 產能上限（全產全銷）與營收")
    row("cx", "平均產值產能指數（FY2026 營收＝1）", "x", {"D": "=1", **{c_(y): f"={P['cx']}_{y % 100}" for y in BU[1:]}}, fmt="0.00",
        note="FY2026 校準：公司稱在產產能滿載，FY2026 營收視為產能（Derived）")
    row("rcap", "產能上限營收 ＝ FY2026 營收 × 產能指數 × 價格指數", "RMB B", {c_(y): f"=a26_rev*{c_(y)}{R['cx']}*{c_(y)}{R['pp']}" for y in BU}, fmt="#,##0.0")
    row("nai", "非 AI 營收", "RMB B", {"C": "=a25_nai", "D": "=a26_nai", **{c_(y): f"={c_(y - 1)}{r}*(1+{P['gn']}{y % 100})" for y in BU[1:]}}, fmt="#,##0.0")
    row("ai", "AI 相關營收 ＝ MIN（需求面、產能上限 − 非 AI）", "RMB B", {"C": "=a25_ai", "D": f"=D{R['aid']}",
        **{c_(y): f"=MIN({c_(y)}{R['aid']},{c_(y)}{R['rcap']}-{c_(y)}{R['nai']})" for y in BU[1:]}}, fmt="#,##0.0", note="產能不足時，非 AI 照成長率出貨、AI 取剩餘產能（Assumed）")
    row("bind", "產能是否受限", "", {c_(y): f'=IF({c_(y)}{R["aid"]}>{c_(y)}{R["rcap"]}-{c_(y)}{R["nai"]}+0.0001,"受限","未受限")' for y in BU[1:]}, fmt="General")
    row("rev", "營收", "RMB B", {c_(y): f"={c_(y)}{R['ai']}+{c_(y)}{R['nai']}" for y in FYS}, fmt="#,##0.0")
    row("g", "營收年增", "%", {c_(y): f"={c_(y)}{R['rev']}/{c_(y - 1)}{R['rev']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("aish", "AI 相關 ÷ 營收", "%", {c_(y): f"={c_(y)}{R['ai']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("util", "需求面 ÷ 產能上限（AI 需求面＋非 AI）", "%", {c_(y): f"=({c_(y)}{R['aid']}+{c_(y)}{R['nai']})/{c_(y)}{R['rcap']}" for y in BU}, fmt="0.0%",
        note="大於 100%＝需求超過產能（營收由產能決定）")
    row("pgwe", "每 GW 勝宏 AI 相關營收", "US$M/GW", {c_(y): f"={c_(y)}{R['ai']}/fx/{c_(y)}{R['gw']}*1000" for y in BU}, fmt="#,##0")
    sec_("C. 損益")
    row("gmn", "其他營收毛利率 ＝ 錨點毛利率 − 差距 × AI 占比（FY2026 校準）", "%", {"D": f"=a26_gp/D{R['rev']}-{P['d']}*D{R['aish']}"}, fmt="0.0%",
        note="Derived：讓 FY2026 毛利率等於錨點；AI 相關毛利率＝其他＋差距")
    row("gma", "AI 相關毛利率 ＝ 1 −（1 − 校準值）÷ 價格指數 ＋ 額外改善", "%", {"D": f"=D{R['gmn']}+{P['d']}", **{c_(y): f"=1-(1-$D${R['gmn'] + 1})/{c_(y)}{R['pp']}+{P['gmu']}_{y % 100}" for y in BU[1:]}}, fmt="0.0%",
        note="Derived：同規格成本不變，價格升降全部落在毛利；額外改善為新廠爬坡完成與組合（Assumed）")
    row("gp", "毛利", "RMB B", {"C": "=a25_gp", "D": "=a26_gp", **{c_(y): f"={c_(y)}{R['ai']}*{c_(y)}{R['gma']}+{c_(y)}{R['nai']}*$D${R['gmn']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("gm", "毛利率", "%", {c_(y): f"={c_(y)}{R['gp']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("opb", "營業費用（基準路徑）", "RMB B", {"C": "=a25_opx", "D": "=a26_opx", **{c_(y): f"={c_(y - 1)}{r}*(1+MAX(0,{P['phi']}*{c_(y)}{R['g']}))" for y in BU[1:]}}, fmt="#,##0.0",
        note="FY2025 含財務費用、FY2026 起不含（口徑見 Q 頁）")
    row("opx", "營業費用", "RMB B", {"C": f"=C{R['opb']}", "D": f"=D{R['opb']}", **{c_(y): f"={c_(y)}{R['opb']}+{P['opxa']}*{c_(y)}{R['rev']}" for y in BU[1:]}}, fmt="#,##0.0")
    row("oi", "營業利益", "RMB B", {c_(y): f"={c_(y)}{R['gp']}-{c_(y)}{R['opx']}" for y in FYS}, fmt="#,##0.0")
    row("oim", "營業利益率", "%", {c_(y): f"={c_(y)}{R['oi']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    row("nop", "營業外與財務淨額", "RMB B", {c_(y): "=nop_f" for y in BU[1:]}, fmt="#,##0.0", note="FY2025、FY2026 含在淨利錨點內")
    row("ni", "歸屬母公司淨利", "RMB B", {"C": "=a25_ni", "D": "=a26_ni", **{c_(y): f"=({c_(y)}{R['oi']}+{c_(y)}{R['nop']})*nir_f" for y in BU[1:]}}, fmt="#,##0.0")
    row("nioi", "淨利 ÷ 營業利益", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['oi']}" for y in FYS}, fmt="0.0%")
    row("sh_", "股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", "E": "=sh_h2*(1+sh_g)", **{c_(y): f"={c_(y - 1)}{r}*(1+sh_g)" for y in (2028, 2029)}}, fmt="0.0000",
        note="FY2027 起＝A＋H 股期末股數（FY2026 為加權）")
    row("eps", "EPS", "RMB", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['sh_']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {c_(y): f"={c_(y)}{R['eps']}/{c_(y - 1)}{R['eps']}-1" for y in FYS[1:]}, fmt="0.0%")
    row("nim", "淨利率（歸屬母公司淨利 ÷ 營收）", "%", {c_(y): f"={c_(y)}{R['ni']}/{c_(y)}{R['rev']}" for y in FYS}, fmt="0.0%")
    # FY2030–32：收斂法（上游只到 2029；非驅動因子推導）
    fg, fm = f"fade_g{s}", f"fade_m{s}"
    for y in FADE:
        c, pc, k = FC[y], FC[y - 1], y - 2029
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+({fg}-$G{R['g']})*{k}/3"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({fm}-$G{R['nim']})*{k}/3"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['sh_']}"] = f"={pc}{R['sh_']}*(1+sh_g)"
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['sh_']}"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fmt_ in (("g", "0.0%"), ("rev", "#,##0.0"), ("nim", "0.0%"), ("ni", "#,##0.0"), ("sh_", "0.0000"), ("eps", "#,##0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fmt_; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 11, "FY2030 起為收斂法：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 勝宏科技損益（十億人民幣；EPS 人民幣；歸屬母公司）。FY2025A 實績、FY2026 錨點（1H26 實績＋3Q、4Q 假設），FY2027 起為模型。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要。FY2026–FY2027 三情境相同（顆數係數從 CY2028 開始）"
PL["A2"] = "灰底＝FY2030–FY2032 收斂法（上游只到 2029）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 J 節），非驅動因子推導"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數基準（係數 1.0）", "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85", "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目"] + YH[1:])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "RMB B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("ai", "  AI 相關營收（AI 算力、資料中心、HPC）", "RMB B", "#,##0.0"),
        ("nai", "  非 AI 營收", "RMB B", "#,##0.0"), ("util", "  需求面 ÷ 產能上限", "%", "0.0%"), ("gm", "毛利率", "%", "0.0%"), ("oim", "營業利益率", "%", "0.0%"),
        ("ni", "歸屬母公司淨利", "RMB B", "#,##0.0"), ("nim", "  淨利率", "%", "0.0%"), ("eps", "EPS", "RMB", "0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg", "util") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{SUMR[('A', 'nim')]}")
IN.cell(INROW["fade_mB"], 5, f"=(PL!G{SUMR[('B', 'nim')]}+nim_avg)/2")
for k in ("fade_mA", "fade_mB"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 錨點").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（RMB B）", "+".join(f"ABS(D{PB[s]['rev']}-a26_rev)" for s in SC)),
               ("  情境 A：AI 相關差（RMB B）", f"D{PB['A']['ai']}-a26_ai"),
               ("  情境 A：毛利差（RMB B）", f"D{PB['A']['gma']}*D{PB['A']['ai']}+D{PB['A']['gmn']}*D{PB['A']['nai']}-a26_gp"),
               ("  情境 A：產能上限 − 營收（RMB B；FY2026 校準恆等）", f"D{PB['A']['rcap']}-D{PB['A']['rev']}")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 10; PL.column_dimensions["K"].width = 70
for c in "CDEFGHIJ": PL.column_dimensions[c].width = 13

# ---------------- Units：每 GW PCB 金額與份額（權威） ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 每 GW AI 高速 PCB 金額（加速器卡／運算板、交換器板、中板／背板／其他）、勝宏份額與產能（本專案權威，規則第 13 條；與台光電模型的每 GW CCL 金額口徑分開）。部署 GW＝AI IT 資本支出（markets/network）÷ 每 GW IT 資本支出（Tokenomics，經橋接表）"
A_ = PB["A"]
cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
for key, lab, unit, f, fmt, note in (
        ("gw", "部署 GW（情境 A）", "GW", lambda y: f"=PL!{FC[y]}{A_['gw']}", "0.0", "AI IT 資本支出 ÷ 每 GW IT 資本支出"),
        ("vcmp", "每 GW PCB：加速器卡／運算板", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['vcmp']}", "#,##0", "2026 Assumed、之後 × 指數（Analogy）"),
        ("vsw", "每 GW PCB：交換器板", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['vsw']}", "#,##0", "Assumed"),
        ("vmid", "每 GW PCB：中板／背板／其他", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['vmid']}", "#,##0", "Assumed；VR200 中板與 CPX、Kyber 正交背板"),
        ("pgw", "每 GW AI 高速 PCB 金額（合計，不含價格指數）", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgw']}", "#,##0", ""),
        ("prk", "每架等值 PCB 金額（以 VR200 每 GW 架數換算）", "US$K/架", lambda y: f"=PL!{FC[y]}{A_['pgw']}*1000/racks_vr", "#,##0.0", "對照 Morgan Stanley VR200 每架 US$116.7K（只含機架內）"),
        ("ccl", "對照：台光電模型每 GW 高速 CCL 金額（CCL 口徑）", "US$M/GW", lambda y: f"=emc_ccl{y % 100}", "#,##0", "唯讀 emc/dist；PCB ÷ CCL 見下一列"),
        ("pcr", "每 GW PCB 金額 ÷ 每 GW CCL 金額", "x", lambda y: f"=PL!{FC[y]}{A_['pgw']}/emc_ccl{y % 100}", "0.00", "PCB 售價含 CCL、銅箔以外的製程價值；兩模型範圍不同（見 Recon）"),
        ("mku", "AI 高速 PCB 市場（含價格指數）", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}*PL!{FC[y]}{A_['pp']}", "#,##0.0", "含 ASIC 部署（假設每 GW 金額相同）"),
        ("sh", "勝宏有效份額", "%", lambda y: f"=PL!{FC[y]}{A_['sh']}", "0.0%", "FY2026 以錨點 AI 相關營收校準"),
        ("aid", "勝宏 AI 相關營收（需求面）", "US$B", lambda y: f"=PL!{FC[y]}{A_['aid']}/fx", "#,##0.00", ""),
        ("ai", "勝宏 AI 相關營收（實際＝MIN 需求面、產能）", "US$B", lambda y: f"=PL!{FC[y]}{A_['ai']}/fx", "#,##0.00", ""),
        ("pgwe", "每 GW 勝宏 AI 相關營收", "US$M/GW", lambda y: f"=PL!{FC[y]}{A_['pgwe']}", "#,##0", ""),
        ("cx", "平均產值產能指數（FY2026＝1）", "x", lambda y: f"=PL!{FC[y]}{A_['cx']}", "0.00", "公司擴產計畫（Assumed）"),
        ("rcap", "產能上限營收（含價格指數）", "RMB B", lambda y: f"=PL!{FC[y]}{A_['rcap']}", "#,##0.0", "FY2026 校準"),
        ("util", "需求面 ÷ 產能上限", "%", lambda y: f"=PL!{FC[y]}{A_['util']}", "0.0%", "大於 100%＝產能受限"),
        ("oth", "其他廠商 AI 高速 PCB（沪电、深南、生益电子、鹏鼎、臺美廠等）＝ 市場 − 勝宏", "US$B", lambda y: f"=PL!{FC[y]}{A_['mku']}*PL!{FC[y]}{A_['pp']}-PL!{FC[y]}{A_['ai']}/fx", "#,##0.0", "產能受限時，未滿足需求計入其他廠商")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit)
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：每 GW PCB 金額的絕對水準只影響 FY2026 校準出的份額（金額高 → 份額低），不影響營收；影響營收的是金額指數（世代升級）與份額相對值。FY2027 需求面高於產能時，營收由產值產能指數決定，擴產進度比需求更關鍵；FY2029 起擴產趕上需求，份額與每 GW 金額重新成為主要變數。").font = BOLD
UNOTE = ur
UN.column_dimensions["A"].width = 66; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 14
UN.column_dimensions["G"].width = 60

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（人民幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」；目標價＝DCF 與 Comparable 加權（第 0 節）；十年 P/E 規則值（第 3 節，年底落後 P/E）只作參考。本版未做先驗，無 Andy 合理 P/E 對照"
V["A2"] = "循環位置：勝宏不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照台光電、Arista、台達電；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 51
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-23 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "0.00"
for k, y, lab in ((3, 2026, "現價 ÷ FY2026 EPS（錨點）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 A）≈ 0.2 × FY2026＋0.8 × FY2027"); V[f"B{O + 7}"] = f"=0.2*PL!D{SUMR[('A', 'eps')]}+0.8*PL!E{SUMR[('A', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 A）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（stockanalysis，10-23）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis，10-23）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 3}"] = "FY2026 三情境相同（錨點）"; V[f"E{O + 4}"] = "FY2027 三情境相同（顆數係數從 CY2028 開始）"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, CY0 + 1, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（人民幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數（參考，不是目標價）").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "0.00"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
RB0 = vr + 3
# 十年 P/E：companiesmarketcap 年底 P/E（2016–2025）；自算＝stockanalysis 年底市值 ÷ 歸母淨利（2021–2025）
PE_CMC = [33.8, 30.9, 21.4, 25.9, 30.4, 34.2, 13.8, 23.4, 34.5, 59.2]
MC = [None] * 5 + [26.13, 11.18, 15.75, 35.93, 250.23]
NIY = [None] * 5 + [0.67042, 0.79065, 0.67135, 1.15443, 4.31199]
PE_SELF = [None if m is None else f"={m}/{n}" for m, n in zip(MC, NIY)]
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2016–2025 共 10 年）"},
    {"label": "自算（stockanalysis 年底市值 ÷ 歸母淨利）", "vals": PE_SELF, "tag": "Derived（2021–2025；2016–2020 年底市值為付費資料：未找到；只作核對，不採用）", "use": False}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（A）", "basis": f"=C{VAL['A']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位。採用一個來源（companiesmarketcap，10 年）；自算只有 2021–2025，與採用值同量級（差 0–4 倍），只作核對。2016–2024 年勝宏以汽車、手機、PC 板為主（14–35x），2025 年 AI 轉型後 59x；十年分布含 A 股結構性溢價。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025", (V.cell(RB0 + 5, 1).value, V.cell(RB0 + 14, 1).value)
PE10 = f"$B${RB0 + 5}:$B${RB0 + 14}"
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# 0. 目標價（Andy 2026-10-25 11:28／11:53：DCF 與 Comparable 加權；列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期）
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "0.00"); EPSR[y] = 8 + i
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
CFR = {}
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)" for c in "BCD"], "0.00"); CFR[y] = 16 + i
def dcf_formula(c, r="tp_r", g="tp_g", reinv="$B$15"):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027。"""
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"({pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6" for c in "BCD"], "0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E；A 股公司只用 A 股同業）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("沪电股份（WUS Printed Circuit）", "002463.SZ", 20.66, "2026-10-22"), ("深南电路（Shennan Circuits）", "002916.SZ", 40.68, "2026-09-14"),
         ("生益电子（Shengyi Electronics）", "688183.SS", 29.59, "2026-09-28"), ("鹏鼎控股（Avary）", "002938.SZ", 31.99, "2026-09-18")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
V.cell(32, 1, "並列參考（非 A 股，不進目標價）：金像電 2368.TW 前瞻 P/E 19.6x（10-22）、TTM Technologies TTMI 10.3x（10-23）；stockanalysis，Interested-party")
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "A 股 PCB 同業：沪电、深南、生益电子、鹏鼎（券商常用比較組；摩根大通稱 A 股 PCB 同業 2027 年平均約 30x）；景旺电子前瞻 P/E 未找到")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(2/12*B{EPSR[2027]}+10/12*B{EPSR[2028]})+tp_wB*(2/12*C{EPSR[2027]}+10/12*C{EPSR[2028]})+tp_wC*(2/12*D{EPSR[2027]}+10/12*D{EPSR[2028]})", None, None], "0.00",
     "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "", "", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "0.00", "stockanalysis：8 位分析師平均 388.30（2026-10-23 擷取）；多數目標價訂於 9–10 月股價下跌前")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-23 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝現價 ÷ 前瞻 P/E（In 頁 cons_feps）")
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
tset(46, "band", "價格帶（DCF 不如預期～順利）", ['=TEXT(D23,"0.00")&"～"&TEXT(B23,"0.00")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}<MIN({PE10}),0,IF(px/{c}{EPSR[2028]}>MAX({PE10}),1,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 6 節 companiesmarketcap）的分位；只作檢查，不進目標價")
V.cell(48, 1, "目標價方法")
V["B48"] = ("DCF：股東現金流＝歸母 EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為近三年加權平均約 20%（再投資率約 20%），FY2027–32 明確期＋永續；"
            "勝宏 2026 年固定資產投資最高 180 億元（約營收的六成），但本模型未建資本支出與折舊，未改用淨利＋折舊－資本支出，擴產期現金流偏高估。"
            "Comparable：沪电、深南、生益电子、鹏鼎前瞻 P/E 中位數（只用 A 股同業；A 股倍數含結構性溢價，且部分價格為 9 月），偏高估。"
            "情境對應：A＝順利（顆數基準）、B＝平平（顆數 ×0.8／0.85）、C＝不如預期（顆數 ×0.6）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；有效份額、其他營收毛利率在每個案例內依 FY2026 錨點重算，FY2026 恆等於錨點）：對 FY2027、FY2028 的營收、AI 相關營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 AI 相關", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 AI 相關", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
def alias(new, base, suf, years=(27, 28, 29), sep="_"):
    for y in years:
        wb.defined_names[f"{new}_{y}"] = DefinedName(f"{new}_{y}", attr_text=wb.defined_names[f"{base}{sep}{y}{suf}"].attr_text)
for nm, base, suf in (("CXLO", "cmpx", "_lo"), ("CXHI", "cmpx", "_hi"), ("SXLO", "swx", "_lo"), ("SXHI", "swx", "_hi"), ("MDLO", "midx", "_lo"), ("MDHI", "midx", "_hi"),
                      ("RLO", "rel", "_lo"), ("RHI", "rel", "_hi"), ("PPLO", "pp", "_lo"), ("PPHI", "pp", "_hi"), ("CAPLO", "cx", "_lo"), ("CAPHI", "cx", "_hi"),
                      ("GULO", "gmu", "_lo"), ("GUHI", "gmu", "_hi")):
    alias(nm, base, suf)
for nm, suf in (("GNLO", "_lo"), ("GNHI", "_hi")):
    for y in (27, 28, 29):
        wb.defined_names[f"{nm}{y}"] = DefinedName(f"{nm}{y}", attr_text=wb.defined_names[f"g_nai{y}{suf}"].attr_text)
cases = [("每 GW 金額成長 低", "三項指數皆取低值（運算板 2027 1.25、交換器 1.15、中板 1.40）", dict(cmpx="CXLO", swx="SXLO", midx="MDLO")),
         ("每 GW 金額成長 高", "三項指數皆取高值（運算板 2027 1.75、交換器 1.60、中板 2.60）", dict(cmpx="CXHI", swx="SXHI", midx="MDHI")),
         ("中板／背板成長 低", "只動中板／背板／其他指數（2027 1.40、2028 1.80）", dict(midx="MDLO")),
         ("中板／背板成長 高", "只動中板／背板／其他指數（2027 2.60、2028 4.00）", dict(midx="MDHI")),
         ("份額流失（低）", "份額相對值 2027 0.70、2028 0.60、2029 0.55（Rubin 份額傳聞情境）", dict(rel="RLO")),
         ("份額維持（高）", "份額相對值 2027–28 1.00、2029 0.95", dict(rel="RHI")),
         ("價格回落快", "價格指數 2027 0.95、2028 0.90、2029 0.85", dict(pp="PPLO")),
         ("價格延續", "價格指數 2027 1.03、2028 1.02、2029 1.00", dict(pp="PPHI")),
         ("擴產延遲", "產能指數 2027 1.45、2028 1.95、2029 2.40", dict(cx="CAPLO")),
         ("擴產提前", "產能指數 2027 1.90、2028 2.80、2029 3.40", dict(cx="CAPHI")),
         ("AI 毛利率未改善", "額外改善 2027–29 皆 −1pt", dict(gmu="GULO")),
         ("AI 毛利率改善 快", "額外改善 2027 +3pt、2028 +4pt、2029 +5pt", dict(gmu="GUHI")),
         ("AI 部署 ×0.8", "FY2027 起 AI IT 資本支出 ×0.8", dict(mf="mf_lo")), ("AI 部署 ×1.2", "×1.2", dict(mf="mf_hi")),
         ("AI 毛利率差距 10pt", "AI 相關毛利率只高 10 個百分點", dict(d="d_ai_lo")), ("AI 毛利率差距 20pt", "高 20 個百分點", dict(d="d_ai_hi")),
         ("非 AI 成長 低", "2027–29 −5%", dict(gn="GNLO")), ("非 AI 成長 高", "2027–29 +5%", dict(gn="GNHI")),
         ("營業費用彈性 1.0", "基準 0.6", dict(phi="phi_hi"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0.0"), ("D", "ai", 2027, "#,##0.0"), ("E", "gm", 2027, "0.0%"), ("F", "eps", 2027, "0.00"),
                           ("H", "rev", 2028, "#,##0.0"), ("I", "ai", 2028, "#,##0.0"), ("J", "gm", 2028, "0.0%"), ("K", "eps", 2028, "0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"G{i}"] = f"=F{i}/F$4-1"; SE[f"L{i}"] = f"=K{i}/K$4-1"
    for c in "GL": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 45; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "顆數情境 B", "CY2028 ×0.8、CY2029 ×0.85", PB["B"], "PL!")
put_case(nn + 1, "顆數情境 C", "CY2028–29 ×0.6", PB["C"], "PL!")
SENS_LAST = nn + 1; SENS_NOTE = nn + 3
SE.cell(nn + 3, 1, "讀法：FY2027 三情境相同；顆數情境 B、C 的差異從 FY2028 起。產能受限的年度，需求面輸入（每 GW 金額、份額、AI 部署）對營收的影響被產能上限截斷，只有產能與價格會改變營收；價格指數同時改變產能上限營收與 AI 毛利率（成本不變），所以對 EPS 的影響大於對營收。")
TPS0 = nn + 5
SE.cell(TPS0, 1, "目標價參數（只改 DCF 參數；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "0.00"
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.4", "", dict(w="0.4")), ("DCF 權重 0.8", "", dict(w="0.8"))]
for k, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
assert TPS_LAST < 45, TPS_LAST
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 56
for c in "CDEFGHIJKL": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 共識、券商估計、研調每架 PCB 金額與台光電模型的每 GW CCL 金額。只作事後對照，不為對上數字而調參數。共識取自同花順（14 家）；券商多為單一來源"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. 共識與券商估計（校準未使用）").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識（同花順）", "=cons_rev26", pl("rev", 2026), "同花順 14 家（2026-10-24）", "共識隱含下半年 200 億、較上半年 +72%；本模型 3Q +20%、4Q +25%", fmt="#,##0.0")
rec("FY2026 EPS vs 共識（同花順）", "=cons_eps26", pl("eps", 2026), "同花順 14 家", "3Q26 季報（10-29）前；共識最新報告 2026-09-07")
rec("FY2026 EPS vs 摩根大通（06-11）", "=jpm_eps26", pl("eps", 2026), "摩根大通 2026-06-11", "")
rec("FY2027 營收 vs 共識（同花順）", "=cons_rev27", pl("rev", 2027), "同花順 14 家", "", fmt="#,##0.0")
rec("FY2027 營收 vs 摩根大通", "=jpm_rev27", pl("rev", 2027), "摩根大通 2026-06-11", "", fmt="#,##0.0")
rec("FY2027 EPS vs 共識（同花順）", "=cons_eps27", pl("eps", 2027), "同花順 14 家", "區間 10.42–17.45")
rec("FY2027 EPS vs 摩根大通", "=jpm_eps27", pl("eps", 2027), "摩根大通 2026-06-11", "")
rec("FY2027 EPS vs 高盛隱含（目標價 550 ÷ 26.3 倍）", "=gs_eps27", pl("eps", 2027), "高盛 2026-09-03（騰訊新聞轉述）", "由目標價推算，非原文數字")
rec("FY2028 營收 vs 共識（同花順）", "=cons_rev28", pl("rev", 2028), "同花順 14 家", "", fmt="#,##0.0")
rec("FY2028 EPS vs 共識（同花順）", "=cons_eps28", pl("eps", 2028), "同花順 14 家", "區間 14.70–26.17")
rec("FY2028 EPS vs 摩根大通", "=jpm_eps28", pl("eps", 2028), "摩根大通 2026-06-11", "")
RC.cell(rr_, 1, "B. 每 GW PCB 金額 vs 台光電每 GW CCL 金額（口徑分開；機制檢查）").font = BOLD; rr_ += 1
for y in (2026, 2027, 2028):
    rec(f"CY{y} 每 GW PCB 金額（本模型）vs 每 GW CCL 金額（台光電模型）", f"=emc_ccl{y % 100}", pl("pgw", y), "emc/dist Units 頁（唯讀）",
        "差距＝PCB ÷ CCL − 1。PCB 售價含 CCL 與銅箔、鑽孔、電鍍、HDI 增層等製程；本模型範圍含光模組板與電源板，台光電範圍含非 AI 高速材料，比例只作量級檢查" if y == 2026 else "—", fmt="#,##0")
rec("CY2026 每 GW PCB ÷ 每 GW CCL（本模型 ÷ 台光電模型）vs Morgan Stanley 每架 PCB ÷ 每架 CCL（GB300）", "=ms_rk_gb/ms_ccl_gb", lambda s: f"=PL!D{PB[s]['pgw']}/emc_ccl26",
    "Morgan Stanley 每架 PCB 35.1K（國金證券轉引）÷ 每架 CCL 6.695K（Deep News 轉述）", "兩模型的比例遠低於每架口徑：台光電每 GW CCL 金額範圍較寬（含非 AI 高速材料）或本模型每 GW PCB 偏低；水準只影響校準份額，不影響營收成長", fmt="0.00")
rec("CY2027 每 GW 金額指數（÷ CY2026）：PCB（本模型）vs CCL（台光電模型）", "=emc_ccl27/emc_ccl26", lambda s: f"=PL!E{PB[s]['pgw']}/PL!D{PB[s]['pgw']}", "emc/dist", "兩者世代升級幅度的對照", fmt="0.00")
rec("CY2027 每 GW 勝宏 AI 營收 vs 每 GW 台光電 AI 營收（US$M）", "=emc_pg27", pl("pgwe", 2027), "emc/dist", "下游 PCB 廠每 GW 營收 ÷ 上游 CCL 廠每 GW 營收", fmt="#,##0")
rec("CY2027 每 GW 金額指數（÷ CY2026）vs Morgan Stanley 每架 PCB 換算", "=(ms_rk_vr*racks_vr)/(ms_rk_gb*racks_gb*(1-vr26)+ms_rk_vr*racks_vr*vr26)",
    lambda s: f"=PL!E{PB[s]['pgw']}/PL!D{PB[s]['pgw']}", "Morgan Stanley 每架 GB300 35.1K、VR200 116.7K 美元（國金證券轉引）× 每 GW 架數", "對照值以 2026 年 Rubin 占比混合", fmt="0.00")
rec("CY2026 每 GW PCB 金額 vs Morgan Stanley 每架金額 × 每 GW 架數（US$M）", "=(ms_rk_gb*racks_gb*(1-vr26)+ms_rk_vr*racks_vr*vr26)/1000", pl("pgw", 2026),
    "同上", "Morgan Stanley 只含 NVL72 機架內 PCB；本模型另含 scale-out 交換器、光模組、ASIC 伺服器", fmt="#,##0")
rec("FY2026 勝宏有效份額 vs 摩根大通 AI／HPC PCB 份額 1H25（Prismark）", "=jpm_sh", lambda s: f"=PL!D{PB[s]['sh0']}", "摩根大通 2026-06-11", "口徑不同：本模型分母只含 AI 部署（含 ASIC）", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "C. 產能與校準").font = BOLD; rr_ += 1
rec("FY2027 需求面 ÷ 產能上限", None, pl("util", 2027), "—", "大於 100%＝產能受限", fmt="0.0%")
rec("FY2028 需求面 ÷ 產能上限", None, pl("util", 2028), "—", "", fmt="0.0%")
rec("其他營收毛利率（Derived）", None, lambda s: f"=PL!D{PB[s]['gmn']}", "FY2026 錨點", "AI 相關毛利率＝此值＋d_ai", fmt="0.0%")
rec("FY2027 毛利率 vs FY2026（百分點）", "=a26_gp/a26_rev", pl("gm", 2027), "錨點", "摩根大通預估 2027 年 41.3%", fmt="0.0%", pt=True)
rec("4Q26 EPS（模型）", None, lambda s: "=Q!E15", "—", "3Q26 季報 2026-10-29 公布", fmt="0.00")
RC.cell(rr_, 1, "D. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以部署路徑取代「AI IT 資本支出 ÷ 每 GW IT 資本支出」"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RECON_LAST = rr_ - 1
RC.column_dimensions["A"].width = 70
for c in "BCDEFGH": RC.column_dimensions[c].width = 12
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司原文＝勝宏科技定期報告與業績交流會，經媒體轉述；券商與研調皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("公司：2026 年半年度報告（深交所 2026-08-28；營收、營業成本、費用、公允價值變動、所得稅、EPS、股數）", H1R),
        ("二手：華爾街見聞 2026-08-27（1H26 營收、歸母、扣非、EPS、ROE、H 股上市）", WS1H),
        ("二手：鳳凰網（1H26 毛利率、公允價值變動、財務費用、交易性金融資產、借款）", IFH),
        ("二手：老虎 2026-04-28（1Q26 營收 55.19 億、歸母 12.88 億）", Q1A),
        ("公司：2026 半年度業績交流會（同花順 2026-08-28；固投、排單到 2027、非 AI 持平、毛利率承壓原因）", CALL),
        ("二手：證券之星 2026-03-14 FY2025 年報簡析（營收、歸母、毛利率、EPS、4Q25）", FY25),
        ("二手：華爾街見聞 FY2025 年報（加權 ROE 35.56%／13.95%、扣非）", WS25),
        ("二手：界面新聞 2026-03（4Q25 低於預期、2026 投資 200 億、2030 千億產值、季毛利率、AI 占比、產銷量）", JM),
        ("公司：2024 年度業績快報（巨潮 2025-03-10；FY2023 加權 ROE 9.35%）", "https://static.cninfo.com.cn/finalpage/2025-03-10/1222758016.PDF"),
        ("券商：摩根大通首次覆蓋（新浪 2026-06-11；2026–28 預估、份額、毛利率、A 股同業倍數）", JPM),
        ("券商：高盛（騰訊新聞 2026-09-03；目標價 550、26.3 倍 2027 P/E）", GS),
        ("研調：國金證券 2026-05（Morgan Stanley 每架 PCB 金額 GB300 3.51 萬、VR200 11.67 萬美元；層數）", GJ),
        ("共識：同花順 F10 盈利預測（14 家，2026-10-24 擷取）", THS),
        ("市場傳聞與股價：新浪 2026-07-21（Rubin 份額傳聞、降價傳聞、4Q25／1Q26 低於預期）", RUM),
        ("產能整理：東方財富財富號 2026-01-05（各廠產能與產值）", CAP),
        ("研報摘要：外刊庫 2026-10-25（同業 1H26 營收成長、每平方米營收）", WK),
        ("資料商：stockanalysis（10-23 收盤、前瞻 P/E、共識目標價、股數；同業前瞻 P/E；年底市值、年度財報）", SA),
        ("資料商：companiesmarketcap（年底 P/E 2016–2025）", "https://companiesmarketcap.com/victory-giant-technology/pe-ratio/"),
        ("券商歷史數字：民生證券 2017-10-27、2020-02-28、2020-08-19；東吴證券 2018-10-28（2016–2019 營收與淨利）", "https://pdf.dfcfw.com/pdf/H3_AP201711201052546282_1.pdf"),
        ("匯率：中國經濟網 2026-10-16（人民幣中間價 6.7220）", FXS),
        ("市場層：markets/network、markets/chip_bridge、tsmc 成品；emc 成品（台光電每 GW CCL，Recon 用）（本 repo，唯讀）", ""),
        ("無法取得：eniu.com 歷史 P/E（權限請求逾時）；pdf.dfcfw.com 2021-04 報告（權限請求逾時）", "未找到"),
        ("待補（人工）：3Q26 季報（2026-10-29）；分產品（高多層、HDI）營收與毛利率；2020 年營收與淨利；產能（萬㎡）與產值的公司原文；資本支出與折舊；前瞻 P/E 十年歷史（Bloomberg）；景旺电子前瞻 P/E", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("勝宏科技（300476.SZ／02476.HK）收支與估值模型 v0.1（AI 半導體 Project：載板與 PCB 次產業；AI 加速卡與伺服器高層數 PCB、HDI）　2026-10-25", True),
 ("版本：v0.1（初版）。每 GW AI 高速 PCB 金額（加速器卡／運算板、交換器板、中板／背板／其他）、勝宏份額與產值產能在本模型 Units 頁建立，為本專案權威資料（規則第 13 條）；與台光電模型的每 GW CCL 金額口徑分開，Recon 頁對照。", False),
 ("模型命題：勝宏的 AI 營收＝部署 GW（AI IT 資本支出 ÷ 每 GW IT 資本支出）× 每 GW AI 高速 PCB 金額（VR200 中板、CPX、Kyber 正交背板使每 GW 金額上升）× 勝宏份額 × 價格指數，但受產值產能（FY2026 營收 × 擴產指數）限制（全產全銷）。非 AI（汽車、手機、PC、消費）依公司說法持平。", False),
 ("核心問題：勝宏 2026 年 180 億元固定資產投資能讓產能在 2027–29 年擴大多少，擴產趕上需求後份額（沪电、深南、鹏鼎追趕；Rubin 份額傳聞）與價格守不守得住；現價約 15 倍前瞻 P/E 隱含什麼路徑。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。預測期 FY2027–FY2029 由驅動因子推導，FY2030–FY2032 為收斂法。金額：十億人民幣；每 GW 金額以美元計、以匯率換算。口徑：歸屬母公司淨利（中國企業會計準則）。A 股與 H 股同股。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("需求面：markets/network 的 AI IT 資本支出 × 情境顆數係數 ÷ 每 GW IT 資本支出（Tokenomics，經橋接表）＝部署 GW；× 每 GW PCB 金額（2026 年三項合計 US$450M，Assumed；之後 × 世代指數，Analogy／Assumed）× 匯率 × 有效份額（FY2026 以錨點 AI 營收校準）× 份額相對值 × 價格指數（Assumed）。", False),
 ("產能面：FY2026 營收（公司稱在產產能滿載）× 產值產能指數（2027 1.70、2028 2.40、2029 3.00；Assumed，依公司各廠計畫與 2030 年千億產值目標）× 價格指數。AI 相關營收＝MIN（需求面、產能上限 − 非 AI）。", False),
 ("毛利率：AI 相關比其他營收高 d_ai 15pt（Assumed），其他毛利率為 FY2026 殘差（Derived）；價格指數變動全部落在 AI 毛利；另加新廠爬坡完成的額外改善（Assumed）。營業費用：彈性 phi 0.6（Assumed）。淨利 ÷ 稅前 0.867（1H26 實績）。", False),
 ("校準：FY2026 ＝ 1H26 實績＋假設（3Q 季增 20%、4Q 季增 25%、毛利率 33%、營業費用率 6.5%）。3Q26 季報預計 2026-10-29 公布，截至 2026-10-25 尚未公布。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2）＋0.5 × Comparable 目標價（A 股同業沪电、深南、生益电子、鹏鼎前瞻 P/E 中位數 × 12 個月後前瞻 EPS）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解。十年 P/E 規則值 × FY2027 EPS 保留為參考（不是目標價）。情境對應：A＝順利、B＝平平、C＝不如預期（原為顆數情境）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司數字取自定期報告（深交所）與媒體轉述；券商與資料商為 Interested-party。", False),
 ("", False),
 ("已知限制：(1) 每 GW PCB 金額的 2026 年水準與世代指數皆為 Assumed／Analogy（Morgan Stanley 每架金額只含 NVL72 機架內）；水準只影響校準份額，指數影響成長；(2) 產值產能指數為 Assumed，公司只給 2030 年千億目標與各廠產值，爬坡速度未揭露；(3) AI 占比 FY2025 取券商轉引的招股書口徑，FY2026 由「非 AI 持平」反推；(4) 未建資本支出與折舊，毛利率的新廠爬坡效果以額外改善近似，DCF 現金流在擴產期偏高估；(5) 人民幣 2026 年升值（中間價 6.72），出口占 76.8%，匯兌與售價影響未建；(6) 十年 P/E 只有一個採用來源，且 2016–2024 年為 AI 前汽車、消費板期；(7) 3Q26 季報（10-29）後更新錨點；R4 產業部署路徑未建。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170

# ---------------- v0.1.1（2026-10-25）：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00／14:17）----------------
# 新列一律放在 Val、README 工作表既有內容之後，不位移既有列；既有數值不變
VER_PREV, VER_NOW = 'v0.1', 'v0.1.1'
OBS = [
 "目標價人民幣 322.84（DCF 202.71、Comparable 442.98 各半），比現價 183.48（2026-10-23）高 76.0%，級距大於 50%；Comparable 單獨為 +141.4%。主因是倍數不是 EPS：本模型前瞻 EPS 14.39 比共識前瞻 EPS 12.36 高 16.4%，但 A 股 PCB 同業（沪电、深南、生益电子、鹏鼎）前瞻 P/E 中位數 30.8 倍、共識目標價 388.30 隱含 31.4 倍，現價只隱含 14.8 倍；DCF 單獨僅 +10.5%，上檔空間主要來自 Comparable，市場對份額（Rubin 份額傳聞）與擴產的折價未計入 Comparable。",
 "份額、每 GW 金額與擴產最敏感：份額流失（2028 年相對值 0.60）使 FY2028 EPS −23.7%、每 GW 金額成長低值 −20.9%、擴產延遲 −20.0%；顆數情境 C 的 FY2028 EPS 10.86 比情境 A 17.78 低 38.9%，DCF 價格帶 98.55～287.37。",
 "上檔空間不靠比共識樂觀：本模型 FY2027 EPS 12.81（三情境相同）比共識 14.77 低 13.3%、情境 A 的 FY2028 EPS 17.78 比共識 22.13 低 19.7%。下一個檢驗點是 3Q26 季報（2026-10-29）。"
]   # 重點觀察（數字取自重算後的 Val／PL／Sens；與 README.md 相同文字）
UPW_ROW = next(r_ for r_ in range(1, V.max_row + 1) if V.cell(r_, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 1; REC_ROW = GRADE_ROW + 1
_u = f"B{UPW_ROW}"
V.cell(GRADE_ROW, 1, "上檔空間級距")
V[f"B{GRADE_ROW}"] = f'=IF({_u}>0.5,"大於 50%",IF({_u}>0.25,"25%–50%",IF({_u}>=0,"0%–25%",IF({_u}>=-0.25,"0%～−25%",IF({_u}>=-0.5,"−25%～−50%","小於 −50%")))))'
V[f"B{GRADE_ROW}"].number_format = "@"
V[f"E{GRADE_ROW}"] = f"以「上檔空間：加權目標價」（第 {UPW_ROW} 列）分六級（Andy 2026-10-25 14:00）"
V.cell(REC_ROW, 1, "投資建議")
V[f"B{REC_ROW}"] = f'=IF({_u}>0.5,"強力買進",IF({_u}>0.25,"買進",IF({_u}>=-0.25,"中立",IF({_u}>=-0.5,"放空","強力放空"))))'
V[f"B{REC_ROW}"].number_format = "@"
V[f"E{REC_ROW}"] = "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空（剛好 +25% 為中立、+50% 為買進、−25% 為中立、−50% 為放空）"
_t = RD["A1"].value
assert VER_PREV in _t and VER_NOW not in _t, _t
RD["A1"] = _t.replace(VER_PREV, VER_NOW, 1)
for r_ in range(2, RD.max_row + 1):
    _v = RD.cell(r_, 1).value
    if isinstance(_v, str) and _v.startswith("版本："):
        RD.cell(r_, 1, "版本：" + VER_NOW + "（2026-10-25：補上檔空間級距、投資建議、重點觀察，數值不變）。前版：" + _v[3:])
        break
rn = RD.max_row + 1
RD.cell(rn, 1, VER_NOW + "：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁最後新增「上檔空間級距」「投資建議」兩列（第 " + str(GRADE_ROW) + "、" + str(REC_ROW) + " 列，依「上檔空間：加權目標價」的固定規則）；重點觀察見下方三列。")
RD.cell(rn + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"')
for k_, o_ in enumerate(OBS, 1):
    RD.cell(rn + 1 + k_, 1, f"重點觀察 {k_}：" + o_)

order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("VG_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE,
               "UNOTE": UNOTE, "T": T, "TPS0": TPS0, "TPS_LAST": TPS_LAST, "PEER_LAST": PEER_LAST, "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0, "block_end": block_end},
              open(os.environ["VG_ROWMAP"], "w"), ensure_ascii=False, indent=1)
