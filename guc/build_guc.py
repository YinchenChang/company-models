# 創意電子（3443.TW）收支與估值模型 v0.1.1（AI 半導體 Project：運算層，ASIC 設計服務＋量產；台積電體系；樣板 mediatek/、結構參考世芯 alchip 分支、mrvl/）
# 用法：python3 guc/build_guc.py raw.xlsx ；再以 LibreOffice headless 重算
# 會計年度＝日曆年，FY(t)＝CY(t)，不平移。金額：十億新台幣（NT$B）；每顆價格：千美元（US$K）；EPS：新台幣（基本）。
# 營收＝NRE（委託設計）＋量產（turnkey：晶圓、封裝、測試轉售，毛利率低）。本模型拆成四塊：
#   (1) Google 自研 Arm CPU（Axion）量產＝Google TPU 顆數（tsmc 成品）× CPU 對 TPU 配比 × 創意份額 × 每顆營收 × k_c × 匯率
#   (2) AI ASIC 量產（其他 ASIC：Maia 等）＝其他 ASIC 顆數（tsmc 成品）× 創意份額 × 每顆核心售價（avgo 成品 core_oth）× k_o × 匯率
#   (3) 車用（含北美車用客戶 AI 晶片）：FY2026 錨點 × 情境倍數
#   (4) 其他（NRE＋BMC、SSD 控制器、加密貨幣等量產）：簡化成長率
import sys, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.dont_write_bytecode = True      # 不在上游資料夾寫入 __pycache__（上游唯讀）
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, os.path.join(ROOT, "markets", "memory"))
import cycle_position as CP
import valuation_rule as VR
import market as MEM

BLUE = Font(color="0000FF"); BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC")
wb = openpyxl.Workbook()
def hdr(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h); c.font = BOLD; c.fill = HDR
SC = ("A", "B", "C")
XP = ("tpu", "oth")
PNAME = {"tpu": "Google TPU", "oth": "其他 ASIC（MTIA、Maia、Ascend 等）"}

# ---------------- 讀入其他模型的成品（唯讀；不修改它們） ----------------
def only(pattern):
    fs = glob.glob(os.path.join(ROOT, pattern))
    assert len(fs) == 1, (pattern, fs)
    return fs[0]
TSMC_F = only("tsmc/dist/*.xlsx"); AVGO_F = only("avgo/dist/*.xlsx"); MRVL_F = only("mrvl/dist/*.xlsx"); MTK_F = only("mediatek/dist/*.xlsx")
rel = lambda f: os.path.relpath(f, ROOT)
TS = {}
for r in openpyxl.load_workbook(TSMC_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith(("u_tpu", "u_oth", "g29_tpu", "g29_oth", "um"))):
        TS[k] = (r[4].value, f"{rel(TSMC_F)}｜In!E{r[0].row}")
for k in XP:
    assert all(f"u_{k}_{y}" in TS for y in (26, 27, 28)) and f"g29_{k}" in TS, k
for s in SC:      # 情境與行業一致：tsmc 顆數係數＝markets/memory SCEN_UNITS（同一字母）
    assert abs(TS[f"um28_{s}"][0] - MEM.SCEN_UNITS[s][0]) < 1e-12 and abs(TS[f"um29_{s}"][0] - MEM.SCEN_UNITS[s][1]) < 1e-12, s
for k in XP:
    UT = [u_ for kk, nm, u_, g, t, sr in MEM.UNITS if kk == k][0]
    for i, y in enumerate((26, 27, 28)):
        assert abs(TS[f"u_{k}_{y}"][0] - UT[i]) < 1e-12, (k, y)      # tsmc 與 markets/memory 同一份顆數
AV = {}
for r in openpyxl.load_workbook(AVGO_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and (str(k).startswith(("core_", "sh_", "u_"))):
        AV[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(AVGO_F)}｜In!E{r[0].row}")
for k in XP:
    for y in (26, 27, 28):
        assert abs(AV[f"u_{k}_{y}"][0][1] - TS[f"u_{k}_{y}"][0]) < 1e-9, (k, y)   # avgo 與 tsmc 同一份顆數
AVU = {}      # avgo Units 頁：Broadcom 顆數（對照）
for r in openpyxl.load_workbook(AVGO_F, data_only=True)["Units"].iter_rows(min_row=4, max_row=20):
    if r[0].value: AVU[str(r[0].value)] = r[0].row
MV = {}
for r in openpyxl.load_workbook(MRVL_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith(("msh_tpu", "msh_oth")):
        MV[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(MRVL_F)}｜In!E{r[0].row}")
assert len(MV) == 6, MV.keys()
MK = {}
for r in openpyxl.load_workbook(MTK_F, data_only=True)["In"].iter_rows(min_row=4):
    k = r[0].value
    if k and str(k).startswith(("mtk_tpu", "mtk_oth")):
        MK[k] = ((r[3].value, r[4].value, r[5].value), f"{rel(MTK_F)}｜In!E{r[0].row}")
assert len(MK) == 6, MK.keys()

# ---------------- In ----------------
IN = wb.active; IN.title = "In"
IN["A1"] = "In — 輸入（藍字；基準欄 E，低 D、高 F）。金額：十億新台幣（NT$B）；每顆價格：千美元（US$K）；顆數：百萬顆；匯率：新台幣／美元。每個輸入附兩個來源（找不到者寫「未找到」）；公司原文算一個"
IN["A2"] = "標記：Verified／Interested-party／Analogy／Assumed／Derived；Analogy、Assumed 給低／高區間，不當點估計（放情境或 Sens）。SemiAnalysis 未使用。公司財報與法說內容多取自媒體轉述（公司自 2026 年起不再給季度區間指引）；公司說法為 Interested-party；券商、資料商、媒體為 Interested-party"
hdr(IN, 3, ["鍵", "項目", "單位", "低", "基準", "高", "標記", "來源1", "來源2／說明"])
rows = []; sec = lambda t: rows.append(("sec", t)); inp = lambda *a: rows.append(a)
UDN2 = "https://udn.com/news/story/7252/9661575"
CT2 = "https://www.ctee.com.tw/news/20260731701051-430201"
UAG = "https://uanalyze.com.tw/articles/3528855768"
UA1 = "https://uanalyze.com.tw/articles/5074646989"
SEP26 = "https://finance.technews.tw/2026/10/05/guc-3443-202609-financial-report/"
UDN10 = "https://money.udn.com/money/amp/story/5607/9800312"
UDN9 = "https://money.udn.com/money/amp/story/5607/9741296"
UDN4 = "https://money.udn.com/money/amp/story/5607/9468975"
UDN1 = "https://money.udn.com/money/story/5607/9296878"
MS5 = "https://technews.tw/2026/05/05/ms-see-guc-2026-q1/"
BOFA5 = "https://technews.tw/2026/05/08/bofa-see-guc-2026/"
SA = "https://stockanalysis.com/quote/tpe/3443/"
SAF = "https://stockanalysis.com/quote/tpe/3443/forecast/"
SAI = "https://stockanalysis.com/quote/tpe/3443/financials/"
SAR = "https://stockanalysis.com/quote/tpe/3443/financials/ratios/"
SAH = "https://stockanalysis.com/quote/tpe/3443/history/"
SDOG = "https://statementdog.com/analysis/3443/eps"
HIS = "https://histock.tw/stock/3443/%E6%AF%8F%E8%82%A1%E7%9B%88%E9%A4%98"
URLNAME = {UDN2: "經濟日報 2026-07-30（2Q26 財報）", CT2: "工商時報 2026-07-31（2Q26 EPS、外資估計）", UAG: "優分析 2026-09-22（法說指引轉述）", UA1: "優分析 2026-01-30（4Q25 法說轉述）",
           SEP26: "科技新報（9 月營收，2026-10-05）", UDN10: "經濟日報 2026-10-07（3Q26 營收、CPU 出貨估計）", UDN9: "經濟日報 2026-09-08（CPU 出貨估計）",
           UDN4: "經濟日報 2026-04-28（本國投顧：2027 營收結構）", UDN1: "經濟日報 2026-01-31（FY2025 營收結構）", MS5: "科技新報 2026-05-05（摩根士丹利）", BOFA5: "科技新報 2026-05-08（美銀）",
           SA: "stockanalysis.com", SAF: "stockanalysis.com 共識頁", SAI: "stockanalysis.com 財報頁", SAR: "stockanalysis.com 比率頁", SAH: "stockanalysis.com 股價歷史", SDOG: "財報狗", HIS: "HiStock"}

sec("A. AI 加速器顆數（tsmc 成品；與 HBM、台積電、Broadcom、Marvell、聯發科模型同一份；情境係數讀台積電模型＝markets/memory SCEN_UNITS，不另訂）")
for k in XP:
    for i, y in enumerate((26, 27, 28)):
        inp(f"u_{k}_{y}", f"{PNAME[k]} 顆數 CY20{y}", "百萬顆", None, TS[f"u_{k}_{y}"][0], None, "Interested-party" if k == "tpu" else "Derived／Assumed",
            TS[f"u_{k}_{y}"][1], "avgo 成品同值、markets/memory UNITS 同值（建置時檢查）" if i == 0 else "同上")
    inp(f"g29_{k}", f"{PNAME[k]} CY2029 顆數成長", "%", None, TS[f"g29_{k}"][0], None, "Assumed", TS[f"g29_{k}"][1], "markets/memory 無 2029；沿用台積電模型")
for s in SC:
    for y in (28, 29):
        inp(f"um{y}_{s}", f"情境 {s}：CY20{y} 顆數 ÷ 基準", "x", None, TS[f"um{y}_{s}"][0], None, "Assumed", TS[f"um{y}_{s}"][1],
            "A＝基準；B＝AI 資本支出消化；C＝明顯下修（與台積電、Broadcom、Marvell、聯發科模型相同；markets/memory SCEN_UNITS，建置時檢查）" if (s == "A" and y == 28) else "同上")

sec("B. Google 自研 Arm CPU（Axion）量產：CPU 顆數＝TPU 顆數 × CPU 對 TPU 配比 × 創意份額（CPU 不占 TPU 加速器份額）")
inp("att26", "CY2026 CPU 對 TPU 配比（CPU 顆數 ÷ TPU 顆數；校準年）", "x", 0.25, 0.325, 0.40, "Interested-party",
    UDN10 + "：2026 年 CPU 出貨約 125 萬顆（法人估計）÷ 市場層 TPU 3.85 百萬顆", "CY2026 只用於校準（k_c 吸收）；公司未揭露客戶與顆數")
ATT = {"A": 0.6, "B": 0.5, "C": 0.4}
for s in SC:
    inp(f"att_{s}", f"CY2027 起 CPU 對 TPU 配比（情境 {s}）", "x", 0.25 if s == "A" else None, ATT[s], 0.75 if s == "A" else None, "Assumed",
        UDN10 + "：新平台 CPU 對 TPU 配置由 1:4 提升到 1:2（法人）" if s == "A" else "同上",
        UDN4 + "：第八代 AI ASIC 伺服器由每顆 x86 CPU 配 4 或 8 顆 ASIC 改為每顆 Arm CPU 配 2 顆（本國投顧）；A 另含一般運算 Axion 執行個體" if s == "A" else "A 最有利、C 最不利")
inp("gsh26", "創意占 Google 自研 CPU 顆數 CY2026–27", "%", 0.8, 1.0, 1.0, "Assumed", "公司：最大 CPU 客戶產能已確保（2026-01-30 法說，工商時報轉述）", "二手：Google Arm CPU 專案為創意（經濟日報 2026-01-31）；客戶名稱公司未揭露")
GSH = {"A": 1.0, "B": 0.9, "C": 0.7}
for s in SC:
    inp(f"gsh28_{s}", f"創意占 Google 自研 CPU 顆數 CY2028 起（情境 {s}；下一代 CPU）", "%", 0.5 if s == "A" else None, GSH[s], 1.0 if s == "A" else None, "Assumed",
        "下一代 CPU 的設計夥伴：未找到" if s == "A" else "同上", "C＝下一代 CPU 部分轉給其他夥伴或客戶自行設計（COT）" if s == "A" else "A 最有利、C 最不利")
inp("pc_cpu", "每顆 CPU 營收（turnkey：晶圓＋封裝測試轉售；未校準）", "US$K", 0.7, 1.0, 1.5, "Assumed", "公司未揭露；3 奈米大型 Arm CPU 晶圓與封裝成本（Assumed）", "校準係數 k_c 以 FY2026 CPU 營收吸收單價誤差")
PFC = {"A": 1.2, "B": 1.1, "C": 1.0}
for s in SC:
    inp(f"pfc_{s}", f"CY2028 起每顆 CPU 營收係數（下一代 CPU；情境 {s}）", "x", 0.8 if s == "A" else None, PFC[s], 1.5 if s == "A" else None, "Assumed",
        "公司：第四季開始出貨下一代 CPU 產品（經濟日報 2026-10-07 轉述法人）" if s == "A" else "同上", "下一代製程與單價：未找到" if s == "A" else "A 最有利、C 最不利")
inp("uf", "創意 CPU 與 AI ASIC 顆數係數（Sens 用；CY2027 起）", "x", 0.8, 1.0, 1.2, "Assumed", "無", "±20%")

sec("C. AI ASIC 量產（其他 ASIC：Maia 等）：每顆規格讀 Broadcom 模型（唯讀）；份額與 Broadcom、Marvell、聯發科合計 ≤ 100%")
lo, b_, hi = AV["core_oth"][0]
inp("core_oth", "其他 ASIC 每顆核心售價（不含 HBM）", "US$K", lo, b_, hi, "Assumed", AV["core_oth"][1], "Broadcom 模型的 ASIC 每顆規格（本專案權威）；創意 turnkey 只含晶圓與封裝轉售，以 k_o 校準")
for y in (26, 27, 28):
    inp(f"av_oth_{y}", f"Broadcom 占其他 ASIC 顆數 CY20{y}（份額加總檢查用）", "%", None, AV[f"sh_oth_{y}"][0][1], None, "Assumed", AV[f"sh_oth_{y}"][1], "Broadcom 模型")
    inp(f"mv_oth_{y}", f"Marvell 占其他 ASIC 顆數 CY20{y}（份額加總檢查用）", "%", None, MV[f"msh_oth_{y}"][0][1], None, "Assumed", MV[f"msh_oth_{y}"][1], "Marvell 模型")
    inp(f"mk_oth_{y}", f"聯發科占其他 ASIC 顆數 CY20{y}（高值；份額加總檢查用）", "%", None, MK[f"mtk_oth_{y}"][0][2], None, "Assumed", MK[f"mtk_oth_{y}"][1].replace("In!E", "In!F"), "聯發科模型高值（基準 0）；世芯模型只有 Trainium 份額，其他 ASIC 為 0")
for y in (26, 27, 28):
    inp(f"av_tpu_{y}", f"Broadcom＋Marvell＋聯發科占 Google TPU 顆數 CY20{y}（基準；對照）", "%", None, AV[f"sh_tpu_{y}"][0][1] + MV[f"msh_tpu_{y}"][0][1] + MK[f"mtk_tpu_{y}"][0][1], None, "Derived",
        "avgo、mrvl、mediatek 成品 In 頁", "創意不占 TPU 加速器份額（CPU 另計）")
GO = {26: {"B": (0.01, 0.03, 0.06)}, 27: {"A": (0.0, 0.06, 0.10), "B": (0.0, 0.04, 0.10), "C": (0.0, 0.02, 0.10)},
      28: {"A": (0.0, 0.08, 0.12), "B": (0.0, 0.05, 0.12), "C": (0.0, 0.02, 0.12)}}
for y, d in GO.items():
    for s, (lo, b_, hi) in d.items():
        key = "guc_oth_26" if y == 26 else f"go{y}_{s}"
        lab = f"創意占其他 ASIC 顆數 CY20{y}" + ("（三情境相同，校準年）" if y == 26 else f"（情境 {s}" + ("；CY2029 沿用）" if y == 28 else "）"))
        inp(key, lab, "%", lo, b_, hi, "Assumed",
            "美銀：Microsoft Maia 200 開始貢獻營收（科技新報 2026-05-08，Interested-party）" if y == 26 else "同上",
            "CY2026 份額只用於校準（k_o 吸收）；公司未揭露客戶" if y == 26 else ("A 最有利、C 最不利；高值受 1 − Broadcom − Marvell − 聯發科高值限制" if s == "A" else "—"))
        cap = 1.0 - AV[f"sh_oth_{y}"][0][1] - MV[f"msh_oth_{y}"][0][1] - MK[f"mtk_oth_{y}"][0][2]
        assert hi <= cap + 1e-9, (y, s, hi, cap)     # Broadcom＋Marvell（基準）＋聯發科（高）＋創意（高）≤ 100%
inp("m_tk", "AI 量產（CPU、AI ASIC turnkey）毛利率", "x", 0.10, 0.13, 0.16, "Assumed",
    "公司：turnkey 毛利率低於其他業務、2Q26 turnkey 占 83%、毛利率 21.6%（經濟日報 2026-07-30）", "Q 頁核對：以 13% 推得其他業務毛利率約 30%（含 NRE）")
GTA = {"A": 0.01, "B": 0.0, "C": -0.02}
for s in SC:
    inp(f"gta_{s}", f"AI 量產毛利率調整（FY2027 起；情境 {s}）", "pt", None, GTA[s], None, "Assumed", "公司：毛利率受 turnkey 比重壓抑、營業利益率維持穩定（優分析 2026-09-22）" if s == "A" else "同上", "A 最有利、C 最不利")

sec("D. 車用（含北美車用客戶 AI 晶片，2H27 起量產）與其他營收（NRE＋BMC、SSD 控制器、加密貨幣等量產）：簡化成長（FY2027 起；情境別）")
AM = {"A": (4.0, 2.0, 0.20), "B": (3.0, 1.6, 0.10), "C": (2.0, 1.2, 0.0)}
for s in SC:
    inp(f"am27_{s}", f"車用營收 FY2027 ÷ FY2026（情境 {s}）", "x", 1.5 if s == "A" else None, AM[s][0], 5.0 if s == "A" else None, "Assumed",
        "公司：車用 2026 年占營收個位數、2027 年雙位數；北美專案金額大，主要在 2H27 後貢獻（優分析 2026-09-22）" if s == "A" else "同上",
        MS5 + "：北美車用 AI 晶片 2H26 認證、明年中量產（Interested-party）" if s == "A" else "—")
    inp(f"am28_{s}", f"車用營收 FY2028 ÷ FY2027（情境 {s}）", "x", None, AM[s][1], None, "Assumed", "2H27 量產的全年效果" if s == "A" else "同上", "—")
    inp(f"ag29_{s}", f"車用營收年增 FY2029（情境 {s}）", "%", None, AM[s][2], None, "Assumed", "同上", "—")
inp("m_auto", "車用量產毛利率", "x", 0.13, 0.16, 0.20, "Assumed", MS5 + "：北美車用 AI 晶片毛利率高於 15%（Interested-party）", "經濟日報 2026-10-07：北美車用專案 2H27 明顯改善獲利結構（法人）")
GOR = {"A": (0.15, 0.10), "B": (0.08, 0.05), "C": (-0.05, 0.0)}
for s in SC:
    inp(f"g_o27_{s}", f"其他營收年增 FY2027（情境 {s}）", "%", -0.15, GOR[s][0], 0.25, "Assumed",
        "公司：FY2026 NRE 年增 15–19%（優分析 2026-09-22）；BMC 占營收高個位數；加密貨幣降為個位數" if s == "A" else "同上",
        "公司：2 奈米 HBM4 PHY、共同封裝光學（COUPE）認證（經濟日報 2026-09-08）" if s == "A" else "—")
    inp(f"g_oL_{s}", f"其他營收年增 FY2028–29（情境 {s}）", "%", -0.15, GOR[s][1], 0.25, "Assumed", "同上", "—")
inp("gmadj", "其他營收毛利率調整（Sens 用；基準 0）", "pt", -0.03, 0.0, 0.03, "Assumed", "無", "其他營收毛利率為 FY2026 殘差（Derived），之後沿用")

sec("E. FY2026 錨點（1H26 實績＋3Q26 營收實績＋公司方向性說法＋4Q26 假設）與營收拆分")
inp("cpu_sh26", "FY2026 Google CPU 量產 ÷ 營收", "%", 0.35, 0.48, 0.55, "Assumed", "公司：CPU 為最大應用、turnkey 主要成長動能（優分析 2026-09-22）",
    MS5 + "：Google CPU 占 1Q26 營收接近 50%、全年超過 50%（Interested-party）；經濟日報 2026-01-31：法人估 30% 以上")
inp("asic_sh26", "FY2026 AI ASIC 量產（其他 ASIC）÷ 營收", "%", 0.01, 0.03, 0.06, "Assumed", BOFA5 + "：Maia 200 開始貢獻（Interested-party）", "公司未揭露；金額：未找到")
inp("auto_sh26", "FY2026 車用 ÷ 營收", "%", 0.03, 0.05, 0.08, "Interested-party", UAG + "：車用 2026 年占營收個位數（公司）", "取中間值（Assumed）")
inp("cpu_sh25", "FY2025 Google CPU 量產 ÷ 營收（對照用）", "%", 0.10, 0.18, 0.25, "Assumed", UDN1 + "：2025 年 AI 相關營收約 25%", "公司：CSP 占 2025 全年營收 25%（優分析 2025-10-07）；CPU 量產自 2H25")
inp("asic_sh25", "FY2025 AI ASIC 量產 ÷ 營收（對照用）", "%", None, 0.05, None, "Assumed", "AI 相關 25% 扣 CPU", "—")
inp("auto_sh25", "FY2025 車用 ÷ 營收（對照用）", "%", None, 0.03, None, "Assumed", "公司：中國、歐洲 ADAS 已量產（2026-01-30 法說）", "—")
inp("fx26", "FY2026 匯率（CPU、AI ASIC 美元營收換算）", "NT$/US$", 30.5, 31.6, 33.0, "Assumed", "與世芯模型 FY2026 假設相同", "實際 FY2026 平均匯率：未找到")
inp("q3_rev", "3Q26 營收（月營收實績；7–9 月合計）", "NT$B", None, "=44.099-11.448-13.897", None, "Verified", SEP26 + "：前 9 月 440.99 億（9 月 71.45 億）", UDN10 + "：3Q26 187.5 億、季增 34.9%")
inp("q3_gm", "3Q26 毛利率（公司：較 2Q26 再下滑）", "%", 0.18, 0.195, 0.21, "Interested-party", UAG + "：第三季毛利率較第二季下滑、全年 20% 以上", UDN10 + "：第三季毛利率略降（法人）")
inp("q3_opx", "3Q26 營業費用（公司：與 2Q26 約持平）", "NT$B", 1.28, 1.35, 1.45, "Interested-party", UAG, "2Q26 1.326（Derived：毛利 − 營業利益）")
inp("q4_g", "4Q26 營收季增", "%", 0.15, 0.35, 0.50, "Assumed", UAG + "：第四季季增幅度可能進一步擴大（公司）", UDN10 + "：法人估季增 35–40%；下一代 CPU 第四季開始出貨")
inp("q4_gm", "4Q26 毛利率", "%", 0.17, 0.185, 0.20, "Assumed", "公司：全年毛利率 20% 以上（優分析 2026-09-22）", "turnkey 比重續升")
inp("q4_opx", "4Q26 營業費用", "NT$B", 1.35, 1.45, 1.60, "Assumed", "2Q26 研發費用 10.5 億、年增 25%（經濟日報 2026-07-30）", "—")
inp("t_h1", "1H26 有效稅率（推估稅前利益用）", "%", None, "=1-5.228/6.049", None, "Derived", SAI + "：近 12 月（至 2026-06）淨利 52.28 億 ÷ 稅前約 60.49 億", "1Q、2Q26 稅前利益與所得稅：未找到")
inp("q_nonop", "3Q、4Q26 營業外淨額（每季）", "NT$B", 0.05, 0.13, 0.20, "Assumed", "1H26 約 0.27（Derived）", "FY2025 約 0.09（stockanalysis 推得）")
inp("q_tax", "3Q、4Q26 有效稅率", "%", 0.12, 0.15, 0.18, "Assumed", "FY2025 約 15.1%（stockanalysis 推得）", "—")
inp("q_sh", "加權股數", "十億股", None, 0.13401, None, "Interested-party", SA + "：流通股數 134.01M（2026-10-08 擷取）", "FY2025 淨利 ÷ EPS ＝ 134.0M（Derived）")

sec("F. FY2027 起：營業費用、營業外、稅、股數、匯率")
inp("phi", "營業費用成長對營收成長的彈性", "x", 0.2, 0.3, 0.6, "Assumed", "公司：營業利益率維持穩定、全年雙位數（優分析 2026-09-22）", "FY2026 營業費用約 +32%、營收約 +100%（Q 頁）")
inp("g_opx_min", "營業費用年增下限", "%", 0.05, 0.08, 0.12, "Assumed", "研發年增 25%（2Q26）；先進封裝 IP、自建運算中心（2026-01-30 法說）", "—")
inp("opx_adj", "營業費用額外占營收比（Sens 用；基準 0）", "pt", -0.02, 0.0, 0.02, "Assumed", "無", "Sens 用 +2 個百分點")
inp("nonop_f", "營業外淨額（每年）", "NT$B", 0.2, 0.5, 0.8, "Assumed", "FY2025 約 0.09、1H26 約 0.27（Derived）", "現金增加帶動利息收入：金額未找到")
inp("tax", "有效稅率（FY2027 起）", "%", 0.12, 0.15, 0.20, "Assumed", "FY2025 約 15.1%（Derived）", "—")
inp("sh_g", "加權股數年變動", "%", -0.005, 0.0, 0.005, "Assumed", "FY2021–25 約 134–135M（stockanalysis 推得）", "—")
inp("fx", "匯率（FY2027 起；美元營收換算）", "NT$/US$", 30.0, 32.0, 34.0, "Assumed", "與聯發科、世芯、台積電模型相同假設", "車用與其他營收以新台幣成長率推，未另換算")

sec("G. 公司說法、共識與券商（只作 Recon 對帳，不作校準）")
inp("co_gm26", "公司：FY2026 毛利率 20% 以上", "%", None, 0.20, None, "Interested-party", UAG, "下限值")
inp("co_opm26", "公司：FY2026 營業利益率雙位數", "%", None, 0.10, None, "Interested-party", UAG, "下限值")
inp("co_auto27", "公司：FY2027 車用占營收雙位數", "%", None, 0.10, None, "Interested-party", UAG, "下限值")
inp("cons_rev26", "共識營收 FY2026", "NT$B", 51.4, 59.5, 65.0, "Interested-party", SAF + "（17 位分析師；2026-09-04 擷取，早於 3Q26 營收公布）", "FY2027 起共識為付費資料：未找到")
inp("cons_eps26", "共識 EPS FY2026", "NT$", 45.73, 51.74, 60.51, "Interested-party", SAF + "（17 位；2026-09-04 擷取）", "優分析資料庫 58.23（2026-09-22）")
inp("va_eps27", "市場共識 FY2027 EPS（Visible Alpha，工商時報轉述）", "NT$", None, 94.24, None, "Interested-party", CT2, "資料日 2026-07-31")
inp("fb_eps26", "外資 FY2026 EPS（工商時報 2026-07-31）", "NT$", None, 52.96, None, "Interested-party", CT2, "目標價 6,500")
inp("fb_eps27", "外資 FY2027 EPS", "NT$", None, 104.14, None, "Interested-party", CT2, "—")
inp("fb_eps28", "外資 FY2028 EPS", "NT$", None, 155.25, None, "Interested-party", CT2, "—")
inp("ua_eps27", "優分析資料庫 FY2027 EPS", "NT$", None, 139.12, None, "Interested-party", UAG, "—")
inp("cpu_u27", "法人：CY2027 CPU 出貨", "百萬顆", 3.5, 4.0, 4.0, "Interested-party", UDN10 + "：約 400 萬顆", UDN9 + "：由 250 萬顆上修至 350 萬顆")
inp("dom_cpu27", "本國投顧：FY2027 CSP CPU 專案 ÷ 營收", "%", None, 0.503, None, "Interested-party", UDN4, "—")
inp("dom_ts27", "本國投顧：FY2027 北美車用 turnkey ÷ 營收", "%", None, 0.349, None, "Interested-party", UDN4, "本模型車用含中國、歐洲 ADAS")

sec("H. 價格與市場倍數")
inp("px", "股價（2026-10-22 收盤）", "NT$", None, 7585.0, None, "Interested-party", SAH + "：10-22 收盤 7,585", SDOG + "：10/22 7,585（−5.19%）；10-23 收盤：未找到")
inp("pe_ttm", "近 12 月 P/E（現價 ÷ 近 12 月 EPS 38.90）", "x", None, "=px/38.90", None, "Derived", SA + "：近 12 月 EPS 38.90", "財報狗近四季 EPS 39.01")
inp("fpe", "前瞻 P/E（stockanalysis，2026-10-08；股價 8,360）", "x", None, 102.85, None, "Interested-party", SA, "—")
inp("cons_feps", "共識前瞻 12 月 EPS（2026-10-08 股價 ÷ 前瞻 P/E）", "NT$", None, "=8360/fpe", None, "Derived", "stockanalysis 前瞻 P/E 102.85（股價 8,360）", "共識前瞻 EPS 的直接數字：未找到")
inp("cons_tp", "市場共識目標價（賣方平均）", "NT$", 2500, 6722.37, 8000, "Interested-party", SA + "：19 位分析師平均 6,722（2026-10-08 擷取）", "共識頁（2026-09-04）：18 位平均 5,729、中位數 5,800、區間 2,500–8,000")
inp("nim5", "近五年平均淨利率（FY2021–25；對照用）", "%", None, "=AVERAGE(1460/15108,3710/24040,3508/26241,3451/25044,3770/34141)", None, "Derived",
    "stockanalysis 財報頁：FY2021–25 營收與淨利", "十年平均（2016–2025）：未找到（2016–2020 未取得）；收斂法依工作單改用最後一年 × 0.8")

sec("I. 目標價與收斂法參數（黃底；Assumed，Andy 可改）")
TP_KEYS = []
for s, w in (("A", 0.3), ("B", 0.5), ("C", 0.2)):
    inp(f"tp_w{s}", f"目標價情境權重 {s}（A 順利／B 平平／C 不如預期）", "%", None, w, None, "Assumed", "Andy 可改（工作單 2026-10-25 預設 A 0.3／B 0.5／C 0.2）", "三者合計應為 100%")
    TP_KEYS.append(f"tp_w{s}")
inp("tp_r", "折現率／要求報酬", "%", None, 0.10, None, "Assumed", "Andy 可改（工作單預設 10%）", "美、台、日、韓同一值；單一客戶集中度高的風險不改折現率，放在情境 C")
inp("tp_g", "長期成長率", "%", None, 0.04, None, "Assumed", "Andy 可改（工作單預設 4%）", "永續成長與 FY2032 收斂目標（情境 B）")
inp("tp_roe", "ROE（近三年平均，FY2023–25）", "%", None, "=AVERAGE(0.3954,0.3290,0.3114)", None, "Derived", "stockanalysis 比率頁：FY2023 39.54%、FY2024 32.90%、FY2025 31.14%", "單一來源；近 12 月 43.09%（同站）")
inp("tp_w_dcf", "目標價中 DCF 的權重；其餘為 Comparable", "%", None, 0.5, None, "Assumed", "Andy 可改（2026-10-25 11:53 定 DCF／Comparable 各半）", "—")
inp("fade_gA", "情境 A：FY2032 營收年增收斂目標（tp_g＋2pt）", "%", None, "=tp_g+0.02", None, "Assumed", "工作單 2026-10-25 11:20 收斂法", "FY2030 起由 FY2029 年增線性收斂")
inp("fade_gB", "情境 B：FY2032 營收年增收斂目標（＝tp_g）", "%", None, "=tp_g", None, "Assumed", "同上", "—")
inp("fade_gC", "情境 C：FY2032 營收年增收斂目標（tp_g−2pt）", "%", None, "=tp_g-0.02", None, "Assumed", "同上", "—")
inp("fade_mA", "情境 A：FY2032 淨利率收斂目標（＝情境 A FY2029 淨利率）", "%", None, 0, None, "Derived", "PL 頁情境 A FY2029（公式於建置時連結）", "—")
inp("fade_mB", "情境 B：FY2032 淨利率收斂目標（FY2029 與「FY2029 × 0.8」的中點）", "%", None, 0, None, "Assumed", "十年平均未找到，依工作單以最後一年 × 0.8 代替", "—")
inp("fade_mC", "情境 C：FY2032 淨利率收斂目標（FY2029 × 0.8）", "%", None, 0, None, "Assumed", "同上", "—")
TP_KEYS += ["tp_r", "tp_g", "tp_roe", "tp_w_dcf", "fade_gA", "fade_gB", "fade_gC", "fade_mA", "fade_mB", "fade_mC"]

r = 4; INROW = {}
for srow in rows:
    if srow[0] == "sec":
        IN.cell(r, 1, srow[1]).font = BOLD; r += 1; continue
    key, item, unit, lo, base, hi, tag, s1, s2 = srow
    IN.cell(r, 1, key); IN.cell(r, 2, item); IN.cell(r, 3, unit)
    for j, v in ((4, lo), (5, base), (6, hi)):
        if v is not None:
            IN.cell(r, j, v)
            if not (isinstance(v, str) and v.startswith("=")): IN.cell(r, j).font = BLUE
    for u_, nm_ in URLNAME.items():          # 網址只列在 Sources 頁；In 頁寫來源名稱
        s1 = s1.replace(u_, nm_); s2 = s2.replace(u_, nm_)
    IN.cell(r, 7, tag); IN.cell(r, 8, s1); IN.cell(r, 9, s2)
    wb.defined_names[key] = DefinedName(key, attr_text=f"In!$E${r}")
    if lo is not None: wb.defined_names[key + "_lo"] = DefinedName(key + "_lo", attr_text=f"In!$D${r}")
    if hi is not None: wb.defined_names[key + "_hi"] = DefinedName(key + "_hi", attr_text=f"In!$F${r}")
    INROW[key] = r; r += 1
for k in TP_KEYS:
    for c in range(1, 10): IN.cell(INROW[k], c).fill = KEY
for col, w in zip("ABCDEFGHI", [12, 58, 10, 9, 11, 9, 18, 70, 70]): IN.column_dimensions[col].width = w
IN.freeze_panes = "D4"
def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ---------------- Q：實績與 FY2026 錨點（NT$B） ----------------
Q = wb.create_sheet("Q")
Q["A1"] = "Q — 實績與 FY2026 錨點（十億新台幣；EPS 為新台幣、基本；股數為十億股）。藍字＝輸入（媒體轉述公司財報、stockanalysis）；黑字＝公式。CPU、AI ASIC、車用、其他的拆分為 Assumed（公司未揭露金額）"
Q["A2"] = "FY2026 錨點＝1Q、2Q26 實績＋3Q26 營收實績（月營收）與公司方向性說法（毛利率較 2Q 下滑、費用持平）＋4Q26 假設；公司自 2026 年起不給季度區間指引。3Q26 財報與法說預定 2026-10-29（stockanalysis），屆時以實績取代 E 欄"
hdr(Q, 3, ["項目", "FY2025A", "1Q26（實績）", "2Q26（實績）", "3Q26（營收實績＋公司說法）", "4Q26（模型）", "FY2026 錨點", "標記", "來源1", "來源2／說明"])
S1 = "FY2025：stockanalysis 財報頁、經濟日報 2026-01-31；1Q26：科技新報 2026-05-05／05-08；2Q26：經濟日報 2026-07-30"
def qrow(r, label, vals, fmt="#,##0.000", tag="Verified", s1=S1, s2="—", ann=None):
    Q.cell(r, 1, label)
    for col, v in vals.items():
        c = Q[f"{col}{r}"]; c.value = v; c.number_format = fmt
        if not (isinstance(v, str) and v.startswith("=")): c.font = BLUE
    if ann != "":
        Q[f"G{r}"] = ann or f"=SUM(C{r}:F{r})"; Q[f"G{r}"].number_format = fmt
    Q.cell(r, 8, tag); Q.cell(r, 9, s1); Q.cell(r, 10, s2)
qrow(4, "營收", {"B": 34.141, "C": 11.448, "D": 13.897, "E": "=q3_rev", "F": "=E4*(1+q4_g)"}, tag="Verified／Assumed（4Q）", s2="3Q26：月營收實績（科技新報）；1H26 253.4 億（經濟日報）")
qrow(5, "  Google CPU 量產（年度拆分）", {"B": "=B4*cpu_sh25"}, ann="=G4*cpu_sh26", tag="Assumed", s1="In 頁 E 節", s2="季度拆分未揭露；只在年度欄拆分")
qrow(6, "  AI ASIC 量產（其他 ASIC）", {"B": "=B4*asic_sh25"}, ann="=G4*asic_sh26", tag="Assumed", s1="In 頁 E 節")
qrow(7, "  車用", {"B": "=B4*auto_sh25"}, ann="=G4*auto_sh26", tag="Interested-party／Assumed", s1="In 頁 E 節", s2="公司：2026 年車用占營收個位數")
qrow(8, "  其他（NRE＋BMC、SSD 控制器、加密貨幣等量產）", {"B": "=B4-B5-B6-B7"}, ann="=G4-G5-G6-G7", tag="Derived", s1="—", s2="FY2025 NRE 約占 34%（經濟日報 2026-01-31）；公司：FY2026 NRE 年增 15–19%")
qrow(9, "毛利", {"B": 8.454, "C": "=25.345*0.241-D9", "D": "=13.897*0.216", "E": "=E4*q3_gm", "F": "=F4*q4_gm"}, tag="Verified（毛利率 × 營收）／Interested-party（3Q）／Assumed（4Q）",
     s2="2Q26 毛利率 21.6%、1H26 24.1%（經濟日報；工商時報 2Q26 為 21.5%）；1Q26＝1H26 − 2Q26")
qrow(10, "毛利率", {c: f"={c}9/{c}4" for c in "BCDEFG"}, fmt="0.00%", ann="", tag="Derived", s1="—", s2="公司：1Q26 約 27.2%、2Q26 21.6%")
qrow(11, "營業費用", {c: f"={c}9-{c}12" for c in "BCD"} | {"E": "=q3_opx", "F": "=q4_opx"}, tag="Derived／Interested-party／Assumed", s2="毛利 − 營業利益")
qrow(12, "營業利益", {"B": 4.350, "C": "=25.345*0.138-D12", "D": 1.676, "E": "=E9-E11", "F": "=F9-F11"}, tag="Verified／Derived", s2="2Q26 16.76 億；1H26 營業利益率 13.8%（經濟日報）")
qrow(13, "營業外淨額", {c: f"={c}14-{c}12" for c in "BCD"} | {"E": "=q_nonop", "F": "=q_nonop"}, tag="Derived／Assumed", s2="稅前 − 營業利益")
qrow(14, "稅前利益", {"B": 4.442, "C": "=C17/(1-t_h1)", "D": "=D17/(1-t_h1)", "E": "=E12+E13", "F": "=F12+F13"}, tag="Derived",
     s2="FY2025：stockanalysis 稅前淨利率推得約 44.42 億；1Q、2Q26＝淨利 ÷（1 − 近 12 月稅率）")
qrow(15, "所得稅", {c: f"={c}14-{c}17" for c in "BCD"} | {"E": "=E14*q_tax", "F": "=F14*q_tax"}, tag="Derived／Assumed", s2="稅前 − 淨利")
qrow(16, "  有效稅率", {c: f"={c}15/{c}14" for c in "BCDEFG"}, fmt="0.0%", ann="", tag="Derived", s1="—")
qrow(17, "歸屬母公司淨利", {"B": 3.770, "C": 1.646, "D": 1.555, "E": "=E14-E15", "F": "=F14-F15"}, tag="Verified／Derived", s2="無重大非控制權益（假設）")
qrow(18, "加權股數（十億股）", {"B": "=3.770/28.13", "C": "=1.646/12.28", "D": "=1.555/11.61", "E": "=q_sh", "F": "=q_sh"}, fmt="0.00000", ann="=AVERAGE(C18:F18)", tag="Derived", s1="淨利 ÷ EPS（公司）")
qrow(19, "EPS（基本）", {c: f"={c}17/{c}18" for c in "BCDEF"}, fmt="#,##0.00", ann="=G17/G18", tag="Derived", s1="淨利 ÷ 股數")
qrow(20, "  公司揭露 EPS", {"B": 28.13, "C": 12.28, "D": 11.61}, fmt="#,##0.00", ann="", tag="Verified", s2="HiStock：FY2025 28.13（stockanalysis 28.01）、1Q26 12.28、2Q26 11.61")
Q.cell(22, 1, "FY2026 錨點的檢查（參考；校準未使用券商數字）").font = BOLD
checks = [("FY2026 營收年增", "=G4/B4-1", "0.0%"), ("3Q26 營收季增", "=E4/D4-1", "0.0%"), ("FY2026 毛利率（公司：20% 以上）", "=G10", "0.0%"),
          ("FY2026 營業利益率（公司：雙位數）", "=G12/G4", "0.0%"), ("FY2026 其他營收年增", "=G8/B8-1", "0.0%"),
          ("核對：FY2026 其他營收毛利率（以 m_tk、m_auto 推得）", "=(G9-m_tk*(G5+G6)-m_auto*G7)/G8", "0.0%"),
          ("檢查：1H26 營收 自算 − 公司 25.34（應約為 0）", "=C4+D4-25.34", "0.000"), ("檢查：1H26 歸屬淨利 自算 − 公司 3.2（應約為 0）", "=C17+D17-3.2", "0.000"),
          ("檢查：FY2025 營業費用 ÷ 營收", "=B11/B4", "0.0%")]
for i, (lab, f, fm) in enumerate(checks, 23):
    Q.cell(i, 1, lab); Q[f"B{i}"] = f; Q[f"B{i}"].number_format = fm
for nm, rr in (("rev", 4), ("cpu", 5), ("asic", 6), ("auto", 7), ("oth", 8), ("gp", 9), ("opx", 11), ("oi", 12), ("nonop", 13), ("tax", 15), ("ni", 17), ("sh", 18), ("eps", 19)):
    defname(f"a25_{nm}", f"Q!$B${rr}"); defname(f"a26_{nm}", f"Q!$G${rr}")
Q.column_dimensions["A"].width = 56
for c in "BCDEFG": Q.column_dimensions[c].width = 14
Q.column_dimensions["H"].width = 30; Q.column_dimensions["I"].width = 50; Q.column_dimensions["J"].width = 70
Q.freeze_panes = "B4"

# ---------------- 引擎：一個情境的完整推導（PL 與 Sens 共用） ----------------
FYS = (2025, 2026, 2027, 2028, 2029); FC = {2025: "C", 2026: "D", 2027: "E", 2028: "F", 2029: "G", 2030: "H", 2031: "I", 2032: "J"}
FADE = (2030, 2031, 2032); FADEFILL = PatternFill("solid", fgColor="E7E6E6")
BU = (2026, 2027, 2028, 2029); F3 = (2027, 2028, 2029)
DEFP = dict(scn="A", uf="1", att="att_{s}", atp="0", gsh="gsh28_{s}", pfc="pfc_{s}", pc="pc_cpu", gos="", gop="0", co="core_oth", mt="m_tk", gta="gta_{s}",
            ma="m_auto", am27="am27_{s}", am28="am28_{s}", ag29="ag29_{s}", fx="fx", go27="g_o27_{s}", goL="g_oL_{s}", gma="0", phi="phi", opxa="0",
            nonop="nonop_f", tax="tax", fg="fade_g{s}", fm="fade_m{s}")
def engine(ws, r0, title, **over):
    P = dict(DEFP); P.update(over); s = P["scn"]
    p = lambda k: P[k].format(s=s)
    ws.cell(r0, 1, title).font = BOLD
    R = {}; r = r0 + 1
    hdr(ws, r, ["項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"]); r += 1
    def row(key, label, unit, fs, fmt="#,##0.00", note=""):
        nonlocal r
        ws.cell(r, 1, label); ws.cell(r, 2, unit)
        for col, f in fs.items():
            ws[f"{col}{r}"] = f; ws[f"{col}{r}"].number_format = fmt
        ws.cell(r, 11, note); R[key] = r; r += 1
    def sec_(t):
        nonlocal r
        ws.cell(r, 1, t).font = BOLD; r += 1
    def uexp(k, y):
        if y <= 2027: return f"u_{k}_{y % 100}"
        if y == 2028: return f"u_{k}_28*um28_{s}"
        return f"u_{k}_28*(1+g29_{k})*um29_{s}"
    sec_("A. Google 自研 Arm CPU 量產（由下而上；D 欄＝FY2026 校準年；FY(t)＝CY(t)）")
    row("u_tpu", "Google TPU 市場顆數（含情境係數）", "百萬顆", {FC[y]: "=" + uexp("tpu", y) for y in BU}, "0.000")
    row("att", "CPU 對 TPU 配比", "x", {FC[y]: ("=att26" if y == 2026 else f"=MAX(0,{p('att')}+{P['atp']})") for y in BU}, "0.000")
    row("gsh", "創意占 Google 自研 CPU", "%", {FC[y]: ("=gsh26" if y <= 2027 else f"=MIN(1,{p('gsh')})") for y in BU}, "0.0%")
    row("cu", "創意 CPU 顆數", "百萬顆", {FC[y]: f"={FC[y]}{R['u_tpu']}*{FC[y]}{R['att']}*{FC[y]}{R['gsh']}" + ("" if y == 2026 else f"*{P['uf']}") for y in BU}, "0.000")
    row("cpx", "每顆 CPU 營收（未校準；CY2028 起 × 下一代係數）", "US$K", {FC[y]: f"={P['pc']}" + (f"*{p('pfc')}" if y >= 2028 else "") for y in BU})
    row("craw", "CPU 營收（未校準）", "US$B", {FC[y]: f"={FC[y]}{R['cu']}*{FC[y]}{R['cpx']}" for y in BU})
    row("kc", "CPU 校準係數 k_c ＝ FY2026 CPU 營收（錨點）÷ 未校準", "x", {"D": f"=a26_cpu/fx26/D{R['craw']}"}, "0.000",
        note="Derived：以 FY2026 錨點的 CPU 拆分校準；FY2027 起沿用。顆數、配比、單價的合併係數")
    row("cusd", "CPU 營收（美元）", "US$B", {FC[y]: f"={FC[y]}{R['craw']}*$D${R['kc']}" for y in BU})
    row("fxr", "匯率", "NT$/US$", {"D": "=fx26", **{FC[y]: f"={P['fx']}" for y in F3}})
    row("cpu", "Google CPU 量產營收", "NT$B", {"C": "=a25_cpu", **{FC[y]: f"={FC[y]}{R['cusd']}*{FC[y]}{R['fxr']}" for y in BU}})
    row("cppu", "每顆創意 CPU 營收（校準後）", "US$K", {FC[y]: f"=IF({FC[y]}{R['cu']}>0,{FC[y]}{R['cusd']}/{FC[y]}{R['cu']},0)" for y in BU})
    sec_("B. AI ASIC 量產（其他 ASIC：Maia 等；由下而上）")
    row("u_oth", "其他 ASIC 市場顆數（含情境係數）", "百萬顆", {FC[y]: "=" + uexp("oth", y) for y in BU}, "0.000")
    row("osh", "創意份額：其他 ASIC", "%", {FC[y]: ("=guc_oth_26" if y == 2026 else f"=MIN(1,MAX(0,go{min(y, 2028) % 100}_{s}{P['gos']}+{P['gop']}))") for y in BU}, "0.0%")
    row("ou", "創意 AI ASIC 顆數", "百萬顆", {FC[y]: f"={FC[y]}{R['u_oth']}*{FC[y]}{R['osh']}" + ("" if y == 2026 else f"*{P['uf']}") for y in BU}, "0.000")
    row("oraw", "AI ASIC 營收（未校準）＝ 顆數 × 每顆核心售價（Broadcom 模型）", "US$B", {FC[y]: f"={FC[y]}{R['ou']}*{P['co']}" for y in BU})
    row("ko", "AI ASIC 校準係數 k_o ＝ FY2026 AI ASIC（錨點）÷ 未校準", "x", {"D": f"=a26_asic/fx26/D{R['oraw']}"}, "0.000",
        note="Derived：k_o < 1＝創意 turnkey 只含晶圓與封裝轉售，低於 Broadcom 的核心售價")
    row("ousd", "AI ASIC 營收（美元）", "US$B", {FC[y]: f"={FC[y]}{R['oraw']}*$D${R['ko']}" for y in BU})
    row("asic", "AI ASIC 量產營收", "NT$B", {"C": "=a25_asic", **{FC[y]: f"={FC[y]}{R['ousd']}*{FC[y]}{R['fxr']}" for y in BU}})
    sec_("C. 車用與其他營收（FY2027 起簡化成長）")
    row("auto", "車用（含北美車用客戶 AI 晶片）", "NT$B", {"C": "=a25_auto", "D": "=a26_auto", "E": f"=D{r}*{p('am27')}", "F": f"=E{r}*{p('am28')}", "G": f"=F{r}*(1+{p('ag29')})"})
    row("oth", "其他（NRE＋BMC、SSD 控制器、加密貨幣等量產）", "NT$B", {"C": "=a25_oth", "D": "=a26_oth", "E": f"=D{r}*(1+{p('go27')})", "F": f"=E{r}*(1+{p('goL')})", "G": f"=F{r}*(1+{p('goL')})"})
    row("rev", "營收", "NT$B", {FC[y]: f"={FC[y]}{R['cpu']}+{FC[y]}{R['asic']}+{FC[y]}{R['auto']}+{FC[y]}{R['oth']}" for y in FYS})
    row("g", "營收年增", "%", {FC[y]: f"={FC[y]}{R['rev']}/{FC[y - 1]}{R['rev']}-1" for y in FYS[1:]}, "0.0%")
    row("csh", "Google CPU ÷ 營收", "%", {FC[y]: f"={FC[y]}{R['cpu']}/{FC[y]}{R['rev']}" for y in FYS}, "0.0%")
    row("ash", "車用 ÷ 營收", "%", {FC[y]: f"={FC[y]}{R['auto']}/{FC[y]}{R['rev']}" for y in FYS}, "0.0%")
    sec_("D. 毛利與損益")
    row("gpa", "AI 量產毛利 ＝（m_tk＋情境調整）×（CPU＋AI ASIC）", "NT$B",
        {"D": f"={P['mt']}*(D{R['cpu']}+D{R['asic']})", **{FC[y]: f"=({P['mt']}+{p('gta')})*({FC[y]}{R['cpu']}+{FC[y]}{R['asic']})" for y in F3}})
    row("gpv", "車用毛利 ＝ m_auto × 車用", "NT$B", {FC[y]: f"={P['ma']}*{FC[y]}{R['auto']}" for y in BU})
    row("gmx", "其他營收毛利率（FY2026 殘差，之後沿用＋調整）", "%", {"C": f"=(a25_gp-m_tk*(C{R['cpu']}+C{R['asic']})-m_auto*C{R['auto']})/C{R['oth']}",
        "D": f"=(a26_gp-D{R['gpa']}-D{R['gpv']})/D{R['oth']}", **{FC[y]: f"=$D{r}+{P['gma']}" for y in F3}}, "0.00%", note="Derived：讓 FY2026 毛利等於錨點（含 NRE，毛利率高）")
    row("gp", "毛利", "NT$B", {"C": "=a25_gp", "D": "=a26_gp", **{FC[y]: f"={FC[y]}{R['gmx']}*{FC[y]}{R['oth']}+{FC[y]}{R['gpa']}+{FC[y]}{R['gpv']}" for y in F3}})
    row("gm", "毛利率", "%", {FC[y]: f"={FC[y]}{R['gp']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("opx", "營業費用", "NT$B", {"C": "=a25_opx", "D": "=a26_opx", **{FC[y]: f"={FC[y - 1]}{r}*(1+MAX(g_opx_min,{P['phi']}*{FC[y]}{R['g']}))+{P['opxa']}*{FC[y]}{R['rev']}" for y in F3}})
    row("oi", "營業利益", "NT$B", {FC[y]: f"={FC[y]}{R['gp']}-{FC[y]}{R['opx']}" for y in FYS})
    row("oim", "營業利益率", "%", {FC[y]: f"={FC[y]}{R['oi']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    row("nonop", "營業外淨額", "NT$B", {"C": "=a25_nonop", "D": "=a26_nonop", **{FC[y]: f"={P['nonop']}" for y in F3}})
    row("pti", "稅前利益", "NT$B", {FC[y]: f"={FC[y]}{R['oi']}+{FC[y]}{R['nonop']}" for y in FYS})
    row("tax", "所得稅", "NT$B", {"C": "=a25_tax", "D": "=a26_tax", **{FC[y]: f"={FC[y]}{R['pti']}*{P['tax']}" for y in F3}})
    row("ni", "歸屬母公司淨利", "NT$B", {FC[y]: f"={FC[y]}{R['pti']}-{FC[y]}{R['tax']}" for y in FYS})
    row("shr", "加權股數", "十億股", {"C": "=a25_sh", "D": "=a26_sh", **{FC[y]: (f"=q_sh*(1+sh_g)" if y == 2027 else f"={FC[y - 1]}{r}*(1+sh_g)") for y in F3}}, "0.00000")
    row("eps", "EPS（基本）", "NT$", {FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['shr']}" for y in FYS})
    row("epsg", "EPS 年增", "%", {FC[y]: f"={FC[y]}{R['eps']}/{FC[y - 1]}{R['eps']}-1" for y in FYS[1:]}, "0.0%")
    row("nim", "淨利率（歸屬母公司淨利 ÷ 營收）", "%", {FC[y]: f"={FC[y]}{R['ni']}/{FC[y]}{R['rev']}" for y in FYS}, "0.00%")
    for y in FADE:      # FY2030–32：收斂法（上游只到 2029；非驅動因子推導）
        c, pc, k = FC[y], FC[y - 1], y - 2029
        ws[f"{c}{R['g']}"] = f"=$G{R['g']}+({p('fg')}-$G{R['g']})*{k}/3"
        ws[f"{c}{R['rev']}"] = f"={pc}{R['rev']}*(1+{c}{R['g']})"
        ws[f"{c}{R['nim']}"] = f"=$G{R['nim']}+({p('fm')}-$G{R['nim']})*{k}/3"
        ws[f"{c}{R['ni']}"] = f"={c}{R['rev']}*{c}{R['nim']}"
        ws[f"{c}{R['shr']}"] = f"={pc}{R['shr']}*(1+sh_g)"
        ws[f"{c}{R['eps']}"] = f"={c}{R['ni']}/{c}{R['shr']}"
        ws[f"{c}{R['epsg']}"] = f"={c}{R['eps']}/{pc}{R['eps']}-1"
        for key, fm in (("g", "0.0%"), ("rev", "#,##0.00"), ("nim", "0.00%"), ("ni", "#,##0.00"), ("shr", "0.00000"), ("eps", "#,##0.00"), ("epsg", "0.0%")):
            ws[f"{c}{R[key]}"].number_format = fm; ws[f"{c}{R[key]}"].fill = FADEFILL
    ws.cell(R["rev"], 11, "FY2030 起為收斂法：營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m（FY2032 達到），非驅動因子推導")
    R["end"] = r
    return R

# ---------------- PL ----------------
PL = wb.create_sheet("PL")
PL["A1"] = "PL — 創意損益（十億新台幣；EPS 新台幣、基本；歸屬母公司）。FY2025A 實績、FY2026 錨點（三情境相同），FY2027–FY2029 由驅動因子推導，FY2030–FY2032 為收斂法。FY(t)＝CY(t)。三情境各一個完整推導區塊（本頁下方），頂端為摘要"
PB = {}; r0 = 60
SCN_T = {"A": "情境 A（順利）：顆數基準（tsmc 情境 A）；CPU 對 TPU 配比 0.6；下一代 CPU 份額 100%、單價 ×1.2；其他 ASIC 份額 6%／8%；車用 ×4.0／×2.0；其他營收 +15%／+10%；量產毛利率 +1pt",
         "B": "情境 B（平平）：CY2028 顆數 ×0.8、CY2029 ×0.85（tsmc 情境 B）；配比 0.5；下一代 CPU 份額 90%、單價 ×1.1；其他 ASIC 4%／5%；車用 ×3.0／×1.6；其他營收 +8%／+5%",
         "C": "情境 C（不如預期）：CY2028–29 顆數 ×0.6（tsmc 情境 C）；配比 0.4；下一代 CPU 份額 70%、單價 ×1.0；其他 ASIC 2%；車用 ×2.0／×1.2；其他營收 −5%／0%；量產毛利率 −2pt"}
for s in SC:
    PB[s] = engine(PL, r0, SCN_T[s], scn=s)
    r0 = PB[s]["end"] + 2
hdr(PL, 3, ["情境／項目", "單位", "FY2025A", "FY2026（錨點）", "FY2027", "FY2028", "FY2029", "FY2030（收斂）", "FY2031（收斂）", "FY2032（收斂）", "說明"])
sr = 4; SUMR = {}
SUMK = (("rev", "營收", "NT$B", "#,##0.0"), ("g", "  營收年增", "%", "0.0%"), ("cpu", "  Google CPU 量產", "NT$B", "#,##0.0"), ("asic", "  AI ASIC 量產（其他 ASIC）", "NT$B", "#,##0.0"),
        ("auto", "  車用", "NT$B", "#,##0.0"), ("oth", "  其他（NRE＋其他量產）", "NT$B", "#,##0.0"),
        ("gm", "毛利率", "%", "0.00%"), ("oim", "營業利益率", "%", "0.00%"), ("ni", "歸屬母公司淨利", "NT$B", "#,##0.0"), ("nim", "  淨利率", "%", "0.00%"),
        ("eps", "EPS", "NT$", "#,##0.00"), ("epsg", "  EPS 年增", "%", "0.0%"))
LONGK = ("rev", "g", "ni", "nim", "eps", "epsg")
for s in SC:
    PL.cell(sr, 1, f"情境 {s}").font = BOLD; sr += 1
    for key, lab, unit, fmt in SUMK:
        PL.cell(sr, 1, lab); PL.cell(sr, 2, unit)
        for y in FYS + (FADE if key in LONGK else ()):
            if key in ("g", "epsg") and y == 2025: continue
            PL[f"{FC[y]}{sr}"] = f"={FC[y]}{PB[s][key]}"; PL[f"{FC[y]}{sr}"].number_format = fmt
            if y in FADE: PL[f"{FC[y]}{sr}"].fill = FADEFILL
        SUMR[(s, key)] = sr; sr += 1
    sr += 1
IN.cell(INROW["fade_mA"], 5, f"=PL!G{SUMR[('A', 'nim')]}")
IN.cell(INROW["fade_mB"], 5, f"=PL!G{SUMR[('B', 'nim')]}*0.9")
IN.cell(INROW["fade_mC"], 5, f"=PL!G{SUMR[('C', 'nim')]}*0.8")
for k in ("fade_mA", "fade_mB", "fade_mC"): IN.cell(INROW[k], 5).font = Font(color="000000")
PL.cell(sr, 1, "核對（應為 0）：FY2026 模型 − 錨點").font = BOLD; sr += 1
CHK = []
for lab, f in (("  三情境：營收差（NT$B）", "+".join(f"ABS(D{PB[s]['rev']}-a26_rev)" for s in SC)),
               ("  三情境：EPS 差（NT$）", "+".join(f"ABS(D{PB[s]['eps']}-a26_eps)" for s in SC)),
               ("  情境 A：CPU 由下而上 − Q 頁 CPU（NT$B）", f"D{PB['A']['cpu']}-a26_cpu")):
    PL.cell(sr, 1, lab); PL[f"D{sr}"] = "=" + f; PL[f"D{sr}"].number_format = "0.0000"; CHK.append(sr); sr += 1
assert sr < 60, sr
PL.column_dimensions["A"].width = 62; PL.column_dimensions["B"].width = 10; PL.column_dimensions["K"].width = 70
for c in "CDEFGHIJ": PL.column_dimensions[c].width = 13
PL.cell(2, 1, "灰底＝FY2030–FY2032 收斂法（上游顆數只到 CY2029）：營收年增線性收斂到 fade_g、淨利率線性收斂到 fade_m（In 頁 I 節），非驅動因子推導")

# ---------------- Units ----------------
UN = wb.create_sheet("Units")
UN["A1"] = "Units — 創意 CPU 與 AI ASIC 量產（情境 A）。顆數：百萬顆（日曆年；TPU、其他 ASIC 顆數與 HBM、台積電、Broadcom、Marvell、聯發科模型同一份）；其他 ASIC 每顆核心售價取自 Broadcom 模型（本專案權威，唯讀）；每顆價：US$K"
A_ = PB["A"]; cyc = lambda y: chr(ord("C") + y - 2026)
hdr(UN, 3, ["項目", "單位", "CY2026", "CY2027", "CY2028", "CY2029", "說明"])
ur = 4; UR = {}
UN.cell(ur, 1, "1. Google 自研 Arm CPU 與其他 ASIC（情境 A；創意營收口徑＝晶圓與封裝轉售，乘校準係數）").font = BOLD; ur += 1
for key, lab, unit, f, fmt, note in (
        ("u", "Google TPU 市場顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['u_tpu']}", "0.00", "tsmc 成品 In 頁"),
        ("att", "CPU 對 TPU 配比", "x", lambda y: f"=PL!{FC[y]}{A_['att']}", "0.00", "In 頁 B 節（CY2026 法人估計；CY2027 起 Assumed）"),
        ("cu", "創意 CPU 顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['cu']}", "0.000", "TPU 顆數 × 配比 × 創意份額"),
        ("cp", "每顆創意 CPU 營收（校準後）", "US$K", lambda y: f"=PL!{FC[y]}{A_['cppu']}", "0.00", "每顆營收 × k_c"),
        ("cusd", "創意 CPU 營收（美元）", "US$B", lambda y: f"=PL!{FC[y]}{A_['cusd']}", "0.00", "—"),
        ("uo", "其他 ASIC 市場顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['u_oth']}", "0.00", "tsmc 成品 In 頁（含 Ascend 等非台積電產品，份額分母偏大）"),
        ("os", "創意份額：其他 ASIC", "%", lambda y: f"=PL!{FC[y]}{A_['osh']}", "0.0%", "In 頁 C 節（Assumed）"),
        ("ou", "創意 AI ASIC 顆數", "百萬顆", lambda y: f"=PL!{FC[y]}{A_['ou']}", "0.000", "市場層顆數 × 份額"),
        ("ousd", "創意 AI ASIC 營收（美元）", "US$B", lambda y: f"=PL!{FC[y]}{A_['ousd']}", "0.00", "顆數 × 每顆核心售價 × k_o")):
    UN.cell(ur, 1, lab); UN.cell(ur, 2, unit); UR[key] = ur
    for y in BU:
        UN[f"{cyc(y)}{ur}"] = f(y); UN[f"{cyc(y)}{ur}"].number_format = fmt
    UN.cell(ur, 7, note); ur += 1
ur += 1
UN.cell(ur, 1, "2. 份額加總檢查（Broadcom＋Marvell＋聯發科＋創意 ≤ 100%；各公司份額讀各自成品 In 頁；世芯只有 Trainium 份額）").font = BOLD; ur += 1
for s in SC:
    UN.cell(ur, 1, f"Broadcom＋Marvell＋聯發科（高）＋創意：其他 ASIC（情境 {s}）"); UN.cell(ur, 2, "%")
    for y in BU:
        yy = min(y, 2028) % 100
        UN[f"{cyc(y)}{ur}"] = f"=av_oth_{yy}+mv_oth_{yy}+mk_oth_{yy}+PL!{FC[y]}{PB[s]['osh']}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
    UN.cell(ur, 7, "其餘為客戶自行設計（COT）或其他夥伴" if s == "A" else "同上"); UR["sum_" + s] = ur; ur += 1
UN.cell(ur, 1, "創意其他 ASIC 份額上限（1 − Broadcom − Marvell − 聯發科高值）"); UN.cell(ur, 2, "%")
for y in BU:
    yy = min(y, 2028) % 100
    UN[f"{cyc(y)}{ur}"] = f"=1-av_oth_{yy}-mv_oth_{yy}-mk_oth_{yy}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
UN.cell(ur, 7, "In 頁份額高值受此上限限制（建置時檢查）"); UR["cap"] = ur; ur += 1
UN.cell(ur, 1, "Broadcom＋Marvell＋聯發科：Google TPU（基準；創意不占）"); UN.cell(ur, 2, "%")
for y in BU:
    UN[f"{cyc(y)}{ur}"] = f"=av_tpu_{min(y, 2028) % 100}"; UN[f"{cyc(y)}{ur}"].number_format = "0.0%"
UN.cell(ur, 7, "創意的 Google 營收為 CPU，不占 TPU 加速器份額"); UR["tpu"] = ur; ur += 2
UN.cell(ur, 1, "3. 校準值（只用 FY2026 已發生的實績與錨點拆分；不用共識、不用 FY2027 起的目標）").font = BOLD; ur += 1
for key, lab, f, fmt, note in (("kc", "CPU 校準係數 k_c", f"=PL!D{A_['kc']}", "0.000", "FY2026 CPU 營收 ÷ 未校準由下而上"),
                               ("ko", "AI ASIC 校準係數 k_o", f"=PL!D{A_['ko']}", "0.000", "FY2026 AI ASIC ÷（顆數 × Broadcom 核心售價）"),
                               ("gmx", "其他營收毛利率（FY2026 殘差）", f"=PL!D{A_['gmx']}", "0.00%", "讓 FY2026 毛利等於錨點")):
    UN.cell(ur, 1, lab); UN[f"C{ur}"] = f; UN[f"C{ur}"].number_format = fmt; UN.cell(ur, 7, note); UR[key] = ur; ur += 1
ur += 1
UN.cell(ur, 1, "讀法：創意的 AI 營收以 Google 自研 CPU 為主，顆數隨 TPU 顆數與 CPU 對 TPU 配比變動，不占 TPU 加速器份額；其他 ASIC 份額受 Broadcom、Marvell、聯發科份額共同限制。").font = BOLD
UN_NOTE = ur
UN.column_dimensions["A"].width = 56; UN.column_dimensions["B"].width = 12
for c in "CDEF": UN.column_dimensions[c].width = 13
UN.column_dimensions["G"].width = 56

# ---------------- Val ----------------
V = wb.create_sheet("Val")
V["A1"] = "Val — P/E 估值（新台幣）。結論用「現價 ÷ FY2027、FY2028 EPS（模型）」；目標價＝DCF 與 Comparable 加權（第 0 節）；十年 P/E 規則值（第 3 節，年底落後 P/E）只作參考"
V["A2"] = "循環位置：創意不是設備股，不套設備市場見頂年；以本模型營收最高的會計年度為見頂年（比照聯發科、世芯；FY(t)＝CY(t)）。營收到 FY2029 仍在增加時，見頂年未知，判為中段"
O = 51
V.cell(O, 1, "1. 價格與市場隱含倍數").font = BOLD
hdr(V, O + 1, ["項目", "A", "B", "C", "說明"])
V.cell(O + 2, 1, "現價（2026-10-22 收盤）"); V[f"B{O + 2}"] = "=px"; V[f"B{O + 2}"].number_format = "#,##0"
for k, y, lab in ((3, 2026, "現價 ÷ FY2026 EPS（錨點）"), (4, 2027, "現價 ÷ FY2027 EPS（模型）"), (5, 2028, "現價 ÷ FY2028 EPS（模型）"), (6, 2029, "現價 ÷ FY2029 EPS（模型）")):
    V.cell(O + k, 1, lab)
    for j, s in enumerate(SC):
        c = "BCD"[j]; V[f"{c}{O + k}"] = f"=px/PL!{FC[y]}{SUMR[(s, 'eps')]}"; V[f"{c}{O + k}"].number_format = "0.0"
V.cell(O + 7, 1, "未來 12 月 EPS（模型 B）≈ 0.2 × FY2026＋0.8 × FY2027"); V[f"B{O + 7}"] = f"=0.2*PL!D{SUMR[('B', 'eps')]}+0.8*PL!E{SUMR[('B', 'eps')]}"; V[f"B{O + 7}"].number_format = "0.00"
V.cell(O + 8, 1, "  現價 ÷ 未來 12 月 EPS（模型 B）"); V[f"B{O + 8}"] = f"=px/B{O + 7}"; V[f"B{O + 8}"].number_format = "0.0"
V.cell(O + 9, 1, "近 12 月 P/E（現價 ÷ 近 12 月 EPS）"); V[f"B{O + 9}"] = "=pe_ttm"; V[f"B{O + 9}"].number_format = "0.0"
V.cell(O + 10, 1, "前瞻 P/E（stockanalysis，2026-10-08）"); V[f"B{O + 10}"] = "=fpe"; V[f"B{O + 10}"].number_format = "0.0"
V[f"E{O + 3}"] = "FY2026 三情境相同（錨點）"; V[f"E{O + 4}"] = "FY2027 三情境差在 CPU 配比、其他 ASIC 份額、車用與其他營收（顆數情境係數從 CY2028 開始）"
CY0 = O + 12
V.cell(CY0, 1, "2. 循環位置（本模型營收最高年；FY(t)＝CY(t)）").font = BOLD
hdr(V, CY0 + 1, ["情境", "營收最高年（FY2026–FY2029）", "FY2027 判定", "規則倍數（P/E）", "見頂年是否未知", "說明"])
CYR = {}
for i, s in enumerate(SC, CY0 + 2):
    rr = f"PL!D{PB[s]['rev']}:G{PB[s]['rev']}"
    V.cell(i, 1, s); CYR[s] = i
    V.cell(i, 2, f"=INDEX({{2026,2027,2028,2029}},MATCH(MAX({rr}),{rr},0))")
    V.cell(i, 3, '=IF(B{0}=2029,"中段",'.format(i) + CP.excel_rule("2027", f"B{i}") + ")")
    V.cell(i, 4, f'=IF(C{i}="見頂",pe_peak,IF(C{i}="谷底",pe_trough,pe_mid))').number_format = "0.0"
    V.cell(i, 5, f'=IF(B{i}=2029,"是：營收到觀測期末仍在增加","否")')
    V.cell(i, 6, "tools/cycle_position.py excel_rule；營收最高年取代設備市場見頂年" if s == "A" else "同上")
R30 = CY0 + 6
V.cell(R30, 1, "3. 規則值隱含股價（新台幣；約 FY2027 年底，未折現）＝ FY2027 EPS × 規則倍數").font = BOLD
hdr(V, R30 + 1, ["情境", "FY2027 EPS", "（同左）", "規則倍數", "隱含股價", "相對現價", "現價 ÷ FY2027 EPS"])
VAL = {}
for j, s in enumerate(SC):
    i = R30 + 2 + j
    V.cell(i, 1, s)
    V.cell(i, 2, f"=PL!E{SUMR[(s, 'eps')]}").number_format = "0.00"
    V.cell(i, 3, f"=B{i}").number_format = "0.00"
    V.cell(i, 4, f"=D{CYR[s]}").number_format = "0.0"
    V.cell(i, 5, f"=C{i}*D{i}").number_format = "#,##0"
    V.cell(i, 6, f"=E{i}/px-1").number_format = "0%"
    V.cell(i, 7, f"=px/C{i}").number_format = "0.0"
    for c in range(1, 8): V.cell(i, c).fill = KEY
    VAL[s] = i
W0 = R30 + 6
V.cell(W0, 1, "4. What must be true：FY2027 年底價 ＝ 現價 ×（1＋報酬）所需的倍數（FY2028 EPS；約 2027 年底的價格看下一年 EPS）").font = BOLD
hdr(V, W0 + 1, ["FY2027 年底價 ÷ 現價", "所需倍數（÷ FY2028 EPS，A）", "所需倍數（÷ FY2028 EPS，B）", "所需倍數（÷ FY2028 EPS，C）", "對照：十年中位數"])
vr = W0 + 2
for x in (1.0, 1.1, 1.21, 1.5, 2.0):
    V.cell(vr, 1, x).font = BLUE; V.cell(vr, 1).number_format = "0.00"
    for j, s in enumerate(SC):
        V.cell(vr, 2 + j, f"=px*A{vr}/PL!F{SUMR[(s, 'eps')]}").number_format = "0.0"
    V.cell(vr, 5, "=pe_mid").number_format = "0.0"
    vr += 1
V.cell(vr, 1, "讀法：1.00＝價格不變、1.21＝一年約 21%。所需倍數低於歷史中位數時，現價相對便宜；倍數是否合理，另看前瞻 P/E 歷史（Bloomberg）。本版未做先驗，無 Andy 合理 P/E 對照。")
WMT_NOTE = vr
# 十年 P/E：companiesmarketcap 年底 P/E；自算＝digrin 12 月未調整股價 ÷ HiStock EPS（2016–2017 EPS 未找到）
PE_CMC = [16.9, 35.3, 24.9, 46.9, 49.9, 51.3, 22.4, 64.9, 50.7, 72.9]
PX_DEC = [78.5, 255.0, 206.0, 241.0, 339.0, 586.0, 641.0, 1740.0, 1360.0, 2125.0]
EPS_H = [None, None, 7.37, 4.73, 6.34, 10.9, 27.69, 26.18, 25.75, 28.13]
PE_SELF = [f"={p_}/{e_}" if e_ else None for p_, e_ in zip(PX_DEC, EPS_H)]
T10 = vr + 2
V.cell(T10, 1, "5. 十年 P/E（2016–2025 年底；兩來源平均，分位排名用；無負值或 >200 需剔除）").font = BOLD
hdr(V, T10 + 1, ["年底", "兩來源平均 P/E", "companiesmarketcap", "自算"])
RB0 = T10 + 16
for i, y in enumerate(range(2016, 2026)):
    rr = T10 + 2 + i
    V.cell(rr, 1, str(y))
    V.cell(rr, 2, f'=IF(COUNT(C{rr}:D{rr})=0,"",AVERAGE(C{rr}:D{rr}))').number_format = "0.00"
    V.cell(rr, 3, f'=IF(B{RB0 + 5 + i}="","",B{RB0 + 5 + i})').number_format = "0.00"
    V.cell(rr, 4, f'=IF(C{RB0 + 5 + i}="","",C{RB0 + 5 + i})').number_format = "0.00"
PE10 = f"$B${T10 + 2}:$B${T10 + 11}"
V.cell(T10 + 12, 1, "年數"); V[f"B{T10 + 12}"] = f"=COUNT({PE10})"
block_end = VR.add_rule_block(V, RB0, [{"id": "P/E", "name": "P/E", "sources": [
    {"label": "companiesmarketcap（年底 P/E）", "vals": PE_CMC, "tag": "Interested-party（資料商；2016–2025 共 10 年）"},
    {"label": "自算（digrin 12 月未調整股價 ÷ HiStock EPS）", "vals": PE_SELF, "tag": "Derived（價格為月資料的 12 月值；EPS：HiStock 2018–2025；2016–2017 EPS 未找到）"}],
    "cases": [
        {"case": "見頂", "pos": "25 分位", "q": 25, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "中段", "pos": "中位數", "q": 50, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"},
        {"case": "谷底", "pos": "75 分位", "q": 75, "basis_label": "FY2027 EPS（B）", "basis": f"=C{VAL['B']}"}]}],
    "0.00", "規則（Andy 2026-10-13）：合理倍數＝2016–2025 年底 P/E 依循環位置取分位，兩個來源分別取分位再平均。stockanalysis 2021–2025 年底 P/E（53.78、23.15、66.47、52.82、75.55）與自算相近，只作核對。十年 P/E 多為 NRE 比重高、營收規模小的時期（中位數約 50 倍），創意轉為 turnkey 量產為主後，倍數未必適用。")
assert V.cell(RB0 + 5, 1).value == "2016" and V.cell(RB0 + 14, 1).value == "2025", (V.cell(RB0 + 5, 1).value, V.cell(RB0 + 14, 1).value)
rule_rows = [rw for rw in range(RB0, block_end) if str(V.cell(rw, 1).value or "").startswith("規則值：")]
assert len(rule_rows) == 3, rule_rows
for nm, rw in zip(("pe_peak", "pe_mid", "pe_trough"), rule_rows):
    defname(nm, f"Val!$C${rw}")

# 0. 目標價（Andy 2026-10-25 11:28／11:53／12:22：DCF 與 Comparable 各半；DCF 負值取 0；列標籤依工作單逐字；B／C／D＝A 順利／B 平平／C 不如預期）
YRS6 = (2027, 2028, 2029, 2030, 2031, 2032)
V.cell(4, 1, "0. 目標價").font = BOLD
V.cell(5, 1, "0.1 DCF（股東權益現金流，FY2027–FY2032 明確期＋FY2032 後永續）").font = BOLD
hdr(V, 6, ["項目", "A（順利）", "B（平平）", "C（不如預期）", "說明"])
T = {}
def tset(r, key, label, fs, fmt, note=""):
    V.cell(r, 1, label); T[key] = r
    for c, f in zip("BCD", fs):
        if f is None: continue
        V[f"{c}{r}"] = f; V[f"{c}{r}"].number_format = fmt
    if note: V[f"E{r}"] = note
V.cell(7, 1, "情境 EPS（FY2027–FY2032）"); V["E7"] = "PL 頁各情境；FY2030 起為收斂法"; T["eps_h"] = 7
EPSR = {}
for i, y in enumerate(YRS6):
    tset(8 + i, f"eps{y}", f"  FY{y}", [f"=PL!{FC[y]}{SUMR[(s, 'eps')]}" for s in SC], "#,##0.00"); EPSR[y] = 8 + i
V.cell(14, 1, "股東現金流＝EPS ×（1 − tp_g ÷ tp_roe）"); T["cf_h"] = 14
tset(15, "reinv", "  再投資率（tp_g ÷ tp_roe，上限 0.8）", ["=MIN(tp_g/tp_roe,0.8)", None, None], "0.0%", "ROE 為近三年平均（In 頁 tp_roe）")
for i, y in enumerate(YRS6):
    tset(16 + i, f"cf{y}", f"  FY{y}", [f"={c}{EPSR[y]}*(1-$B$15)" for c in "BCD"], "#,##0.00")
def dcf_formula(c, r="tp_r", g="tp_g", reinv="$B$15"):
    """c＝情境欄；r、g 可替換（Sens 用）。基準日 2026 年底，t＝1 對 FY2027；負值取 0（Andy 12:22）。"""
    cf = lambda y: f"Val!{c}{EPSR[y]}*(1-{reinv})"
    pv = "+".join(f"{cf(y)}/(1+{r})^{t}" for t, y in enumerate(YRS6, 1))
    return f"MAX(0,{pv}+{cf(2032)}*(1+{g})/({r}-{g})/(1+{r})^6)"
tset(22, "dcfv", "DCF 情境每股價值（折現率 tp_r）", [f"=MAX(0,SUMPRODUCT({c}16:{c}21,1/(1+tp_r)^{{1;2;3;4;5;6}})+{c}21*(1+tp_g)/(tp_r-tp_g)/(1+tp_r)^6)" for c in "BCD"], "#,##0.00",
     "Σ 明確期現金流 ÷ (1＋r)^t ＋ 永續價值 ÷ (1＋r)^6；永續價值＝FY2032 現金流 ×(1＋g)÷(r−g)；基準日 2026 年底；負值取 0")
tset(23, "dcf12", "  DCF 情境價值推到 12 個月後（×(1＋tp_r)）", [f"={c}22*(1+tp_r)" for c in "BCD"], "#,##0.00")
tset(24, "dcftp", "DCF 目標價（情境加權）", ["=tp_wA*B23+tp_wB*C23+tp_wC*D23", None, None], "#,##0.00", "Σ 權重 × 情境價值 ×(1＋r)；權重見 In 頁")
V.cell(26, 1, "0.2 Comparable（市場主流同業前瞻 P/E）").font = BOLD
hdr(V, 27, ["同業與前瞻 P/E", "代號", "前瞻 12 月 P/E", "日期", "來源", "標記"])
PEERS = [("世芯-KY（Alchip）", "3661.TW", 17.15, "2026-10-22"), ("聯發科（MediaTek）", "2454.TW", 38.75, "2026-10-22"),
         ("博通（Broadcom）", "AVGO", 18.52, "2026-10-23"), ("邁威爾（Marvell）", "MRVL", 40.44, "2026-10-23")]
for i, (nm, tk, pe, dt) in enumerate(PEERS, 28):
    V.cell(i, 1, nm); V.cell(i, 2, tk); c = V.cell(i, 3, pe); c.font = BLUE; c.number_format = "0.0"
    V.cell(i, 4, dt); V.cell(i, 5, "stockanalysis.com forward PE（分析師共識；價格日期如左；與聯發科、世芯模型同日擷取）"); V.cell(i, 6, "Interested-party")
PEER_LAST = 27 + len(PEERS)
tset(33, "pmed", "同業前瞻 P/E 中位數", [f"=MEDIAN(C28:C{PEER_LAST})", None, None], "0.0", "ASIC 設計服務與 ASIC 對手（世芯、聯發科、博通、邁威爾）；不含 A 股（非 A 股公司，Andy 2026-10-25 11:53）；智原資料日 2026-09-04 過舊，未列")
tset(34, "feps", "12 個月後的前瞻 EPS（情境加權）", [f"=tp_wA*(2/12*B{EPSR[2027]}+10/12*B{EPSR[2028]})+tp_wB*(2/12*C{EPSR[2027]}+10/12*C{EPSR[2028]})+tp_wC*(2/12*D{EPSR[2027]}+10/12*D{EPSR[2028]})", None, None], "#,##0.00",
     "約 2027-10 時的未來 12 月 EPS ≈ 2／12 × FY2027＋10／12 × FY2028（會計年度＝日曆年）")
tset(35, "comptp", "Comparable 目標價", ["=B33*B34", None, None], "#,##0.00", "同業中位數 × 前瞻 EPS")
V.cell(37, 1, "0.3 結論").font = BOLD
hdr(V, 38, ["項目", "數值", "", "", "說明"])
tset(39, "tpw", "目標價（DCF 與 Comparable 加權）", ["=tp_w_dcf*B24+(1-tp_w_dcf)*B35", None, None], "#,##0.00", "tp_w_dcf × DCF＋（1 − tp_w_dcf）× Comparable；tp_w_dcf 見 In 頁")
tset(40, "comp2", "Comparable 目標價（市場主流同業法）", ["=B35", None, None], "#,##0.00")
tset(41, "cons", "市場共識目標價", ["=cons_tp", None, None], "#,##0.00", "stockanalysis：19 位分析師平均 6,722，2026-10-08 擷取")
tset(42, "upw", "上檔空間：加權目標價", ["=B39/px-1", None, None], "0.0%", "現價＝In 頁 px（2026-10-22 收盤）")
tset(43, "upc", "上檔空間：Comparable 目標價", ["=B40/px-1", None, None], "0.0%")
tset(44, "geps", "與市場共識目標價的差距：EPS 因素", ["=B34/cons_feps-1", None, None], "0.0%", "本模型前瞻 EPS ÷ 共識前瞻 EPS − 1；共識前瞻 EPS＝2026-10-08 股價 ÷ 前瞻 P/E（In 頁 cons_feps）")
tset(45, "gmul", "與市場共識目標價的差距：倍數因素", ["=(cons_tp/cons_feps)/B33-1", None, None], "0.0%")
V["E45"] = '="共識目標價隱含 "&TEXT(cons_tp/cons_feps,"0.0")&"x，同業中位數 "&TEXT(B33,"0.0")&"x"'
tset(46, "band", "價格帶（DCF 不如預期～順利）", ['=TEXT(D23,"#,##0")&"～"&TEXT(B23,"#,##0")', None, None], "@", "C 情境價值～A 情境價值（推到 12 個月後）")
tset(47, "pr", "歷史倍數檢查：現價隱含 P/E 的十年分位（A／B／C）", [f"=IF(px/{c}{EPSR[2028]}>=MAX({PE10}),1,IF(px/{c}{EPSR[2028]}<=MIN({PE10}),0,PERCENTRANK({PE10},px/{c}{EPSR[2028]})))" for c in "BCD"], "0%",
     "現價 ÷ 各情境 FY2028 EPS 在 2016–2025 年底 P/E（第 5 節兩來源平均）的分位；高於十年最大值記為 100%、低於最小值記為 0%；只作檢查，不進目標價")
V.cell(48, 1, "目標價方法")
V["B48"] = ("DCF：股東現金流＝EPS ×（1 − g／ROE），r＝10%、g＝4%、ROE 為近三年平均約 35%（再投資率約 12%），FY2027–32 明確期＋永續；創意為輕資產設計服務公司（量產為轉售，資本支出小），不改用淨利＋折舊－資本支出；情境價值負值取 0。"
            "Comparable：世芯、聯發科、博通、邁威爾前瞻 P/E 中位數；創意本身前瞻 P/E 約 90–100 倍，遠高於同業，同業法偏向低估市場對創意的定價。"
            "主要偏誤：CPU 顆數綁 TPU 顆數與配比，未含一般運算 Axion 的獨立成長；車用、其他營收為簡化成長率；FY2030 起只用收斂法，DCF 偏低估；單一客戶集中（Google CPU 約一半營收）的下行只反映在情境 C。"
            "情境對應：A＝順利（tsmc 顆數 A、配比 0.6）、B＝平平（顆數 B、配比 0.5）、C＝不如預期（顆數 C、配比 0.4、下一代 CPU 份額 70%）")
for r_ in (24, 35, 39):
    V[f"B{r_}"].fill = KEY
V.column_dimensions["A"].width = 58
for c in "BCDEFGH": V.column_dimensions[c].width = 16

# ---------------- Sens ----------------
SE = wb.create_sheet("Sens")
SE["A1"] = "Sens — 一次動一個輸入（情境 A；FY2026 恆等於錨點，k_c、k_o 與其他營收毛利率在每個案例內重算）：對 FY2027、FY2028 營收、毛利率與 EPS 的影響。每個案例是一個完整推導區塊（本頁下方）"
hdr(SE, 3, ["案例", "設定", "FY2027 營收", "FY2027 毛利率", "FY2027 EPS", "FY2027 EPS 相對基準", "FY2028 營收", "FY2028 毛利率", "FY2028 EPS", "FY2028 EPS 相對基準"])
cases = [("CPU 對 TPU 配比 −0.2", "CY2027 起配比 0.6 → 0.4", dict(atp="(-0.2)")), ("CPU 對 TPU 配比 +0.15", "CY2027 起 0.75（含一般運算 Axion）", dict(atp="0.15")),
         ("下一代 CPU 份額 50%", "CY2028 起（基準 100%）", dict(gsh="gsh28_A_lo")),
         ("CPU 與 ASIC 顆數 ×0.8", "CY2027 起", dict(uf="uf_lo")), ("CPU 與 ASIC 顆數 ×1.2", "CY2027 起", dict(uf="uf_hi")),
         ("下一代 CPU 單價 ×1.5", "CY2028 起（基準 ×1.2）", dict(pfc="pfc_A_hi")), ("下一代 CPU 單價 ×0.8", "CY2028 起", dict(pfc="pfc_A_lo")),
         ("其他 ASIC 份額 +4pt", "CY2027 起（Maia 等放量）", dict(gop="0.04")), ("其他 ASIC 份額 0", "CY2027 起", dict(gop="(-1)")),
         ("量產毛利率 10%", "基準 13%（＋情境 A 調整 1pt）", dict(mt="m_tk_lo")), ("量產毛利率 16%", "", dict(mt="m_tk_hi")),
         ("車用 情境 C 路徑", "FY2027 ×2.0、FY2028 ×1.2", dict(am27="am27_C", am28="am28_C", ag29="ag29_C")),
         ("其他營收 情境 C 路徑", "FY2027 −5%、之後 0%", dict(go27="g_o27_C", goL="g_oL_C")),
         ("其他營收毛利率 −3pt", "FY2027 起", dict(gma="gmadj_lo")), ("營業費用彈性 0.6", "基準 0.3", dict(phi="phi_hi")),
         ("營業費用率 +2pt", "營業費用額外 +2% × 營收", dict(opxa="opx_adj_hi")), ("匯率 30", "基準 32（CPU、ASIC 美元營收換算）", dict(fx="fx_lo")),
         ("稅率 20%", "基準 15%", dict(tax="tax_hi"))]
def put_case(i, name, setting, R, pre):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    for c, key, y, fmt in (("C", "rev", 2027, "#,##0"), ("D", "gm", 2027, "0.00%"), ("E", "eps", 2027, "#,##0.00"),
                           ("G", "rev", 2028, "#,##0"), ("H", "gm", 2028, "0.00%"), ("I", "eps", 2028, "#,##0.00")):
        SE[f"{c}{i}"] = f"={pre}{FC[y]}{R[key]}"; SE[f"{c}{i}"].number_format = fmt
    SE[f"F{i}"] = f"=E{i}/E$4-1"; SE[f"J{i}"] = f"=I{i}/I$4-1"
    for c in "FJ": SE[f"{c}{i}"].number_format = "0.0%"
put_case(4, "基準（情境 A）", "PL 區塊 A", PB["A"], "PL!")
base_row = 40; SB = {}
for nm, setting, ov in cases:
    SB[nm] = engine(SE, base_row, f"案例區塊：{nm}（情境 A）", scn="A", **ov)
    base_row = SB[nm]["end"] + 2
for n, (nm, setting, ov) in enumerate(cases):
    put_case(5 + n, nm, setting, SB[nm], "")
nn = 5 + len(cases)
put_case(nn, "情境 B", "顆數 B；配比 0.5；下一代 CPU 份額 90%；其他營收 +8%", PB["B"], "PL!")
put_case(nn + 1, "情境 C", "顆數 C；配比 0.4；下一代 CPU 份額 70%；量產毛利率 −2pt", PB["C"], "PL!")
SENS_LAST = nn + 1
SE.cell(nn + 3, 1, "讀法：FY2027 最大的單一變數是 CPU 對 TPU 配比（CPU 顆數）與量產毛利率；FY2028 再加上 TPU 顆數情境與下一代 CPU 份額。車用與其他營收的情境差異較小。")
SENS_NOTE = nn + 3
TPS0 = nn + 5
SE.cell(TPS0, 1, "目標價參數（只改 DCF 參數；PL 收斂法的 fade_g 不隨 tp_g 變動）").font = BOLD
hdr(SE, TPS0 + 1, ["案例", "設定", "DCF 目標價", "Comparable 目標價", "加權目標價", "相對基準"])
def tp_case(i, name, setting, r="tp_r", g="tp_g", w="tp_w_dcf"):
    SE.cell(i, 1, name); SE.cell(i, 2, setting)
    reinv = f"MIN({g}/tp_roe,0.8)"
    SE[f"C{i}"] = "=(" + "+".join(f"tp_w{s}*{dcf_formula(c, r, g, reinv)}" for s, c in zip(SC, "BCD")) + f")*(1+{r})"
    SE[f"D{i}"] = f"=Val!B{T['comptp']}"
    SE[f"E{i}"] = f"={w}*C{i}+(1-{w})*D{i}"
    SE[f"F{i}"] = f"=E{i}/E${TPS0 + 2}-1"
    for c in "CDE": SE[f"{c}{i}"].number_format = "#,##0.00"
    SE[f"F{i}"].number_format = "0.0%"
TPCASES = [("基準", "r 10%、g 4%、DCF 權重 0.5", {}), ("折現率 +1pt", "r 11%", dict(r="(tp_r+0.01)")), ("折現率 −1pt", "r 9%", dict(r="(tp_r-0.01)")),
           ("長期成長 +1pt", "g 5%", dict(g="(tp_g+0.01)")), ("長期成長 −1pt", "g 3%", dict(g="(tp_g-0.01)")),
           ("DCF 權重 0.3", "", dict(w="0.3")), ("DCF 權重 0.7", "", dict(w="0.7"))]
for k, (nm, st, kw) in enumerate(TPCASES):
    tp_case(TPS0 + 2 + k, nm, st, **kw)
TPS_LAST = TPS0 + 1 + len(TPCASES)
assert TPS_LAST < 40, TPS_LAST
SE.column_dimensions["A"].width = 30; SE.column_dimensions["B"].width = 44
for c in "CDEFGHIJ": SE.column_dimensions[c].width = 14

# ---------------- Recon ----------------
RC = wb.create_sheet("Recon")
RC["A1"] = "Recon — 本模型 vs 公司說法、共識、券商。只作事後對照，不為對上數字而調參數。券商 FY2027–28 估計與法人 CPU 出貨估計為 Interested-party，不用於校準"
hdr(RC, 3, ["項目", "對照值", "模型 A", "差距 A", "模型 B", "差距 B", "模型 C", "差距 C", "對照來源", "說明"])
rr_ = 4
def rec(label, ref, getter, src, note="", fmt="#,##0.00", pt=False, gap=True):
    global rr_
    RC.cell(rr_, 1, label)
    if ref is not None: RC.cell(rr_, 2, ref).number_format = fmt
    for j, s in enumerate(SC):
        cv, cd = 3 + 2 * j, 4 + 2 * j
        RC.cell(rr_, cv, getter(s)).number_format = fmt
        if ref is not None and gap:
            L = chr(64 + cv)
            RC.cell(rr_, cd, f"={L}{rr_}-B{rr_}" if pt else f"={L}{rr_}/B{rr_}-1").number_format = "0.0%"
    RC.cell(rr_, 9, src); RC.cell(rr_, 10, note); rr_ += 1
pl = lambda key, y: (lambda s: f"=PL!{FC[y]}{PB[s][key]}")
RC.cell(rr_, 1, "A. FY2026 錨點 vs 公司說法與共識").font = BOLD; rr_ += 1
rec("FY2026 營收 vs 共識", "=cons_rev26", pl("rev", 2026), "stockanalysis（17 位；2026-09-04）", "共識資料日在 3Q26 營收公布前；3Q26 營收季增 34.9%，遠高於公司說的高個位數至低雙位數", fmt="#,##0.0")
rec("FY2026 EPS vs 共識", "=cons_eps26", pl("eps", 2026), "stockanalysis（2026-09-04）", "共識區間 45.73–60.51；外資 52.96；優分析 58.23")
rec("FY2026 毛利率 vs 公司 20% 以上（百分點）", "=co_gm26", pl("gm", 2026), "優分析 2026-09-22（法說）", "", fmt="0.0%", pt=True)
rec("FY2026 營業利益率 vs 公司雙位數下限（百分點）", "=co_opm26", pl("oim", 2026), "優分析 2026-09-22（法說）", "", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "B. CPU 與車用 vs 法人（市場層 TPU 顆數 × 配比口徑）").font = BOLD; rr_ += 1
rec("CY2027 創意 CPU 顆數 vs 法人約 400 萬顆", "=cpu_u27", pl("cu", 2027), "經濟日報 2026-10-07", "法人另一說 350 萬顆（2026-09-08）；達到 400 萬顆需配比約 0.73（含一般運算 Axion）", fmt="0.00")
rec("CY2027 達到法人 CPU 顆數所需配比（＝400 萬 ÷ TPU 顆數）", None, lambda s: f"=cpu_u27/PL!E{PB[s]['u_tpu']}", "推導", "市場層 CY2027 TPU 5.5 百萬顆", fmt="0.00")
rec("FY2027 CPU ÷ 營收 vs 本國投顧 50.3%（百分點）", "=dom_cpu27", pl("csh", 2027), "經濟日報 2026-04-28", "", fmt="0.0%", pt=True)
rec("FY2027 車用 ÷ 營收 vs 公司雙位數下限（百分點）", "=co_auto27", pl("ash", 2027), "優分析 2026-09-22（法說）", "", fmt="0.0%", pt=True)
rec("FY2027 車用 ÷ 營收 vs 本國投顧北美車用 34.9%（百分點）", "=dom_ts27", pl("ash", 2027), "經濟日報 2026-04-28", "本國投顧比重遠高於公司說法；北美專案主要在 2H27 後貢獻", fmt="0.0%", pt=True)
rec("FY2027 每顆創意 CPU 營收（校準後）", None, pl("cppu", 2027), "推導", "US$K；k_c × 每顆營收", fmt="0.00")
RC.cell(rr_, 1, "C. EPS vs 券商與共識（校準未使用）").font = BOLD; rr_ += 1
rec("FY2027 EPS vs 市場共識（Visible Alpha）", "=va_eps27", pl("eps", 2027), "工商時報 2026-07-31", "資料日在 3Q26 營收公布前")
rec("FY2027 EPS vs 外資", "=fb_eps27", pl("eps", 2027), "工商時報 2026-07-31", "")
rec("FY2027 EPS vs 優分析資料庫", "=ua_eps27", pl("eps", 2027), "優分析 2026-09-22", "")
rec("FY2028 EPS vs 外資", "=fb_eps28", pl("eps", 2028), "工商時報 2026-07-31", "")
rec("共識前瞻 12 月 EPS vs 模型（情境加權 2／12 × FY2027＋10／12 × FY2028）", "=cons_feps", lambda s: f"=Val!B{T['feps']}", "stockanalysis 前瞻 P/E 推得", "三欄同值（情境加權）")
RC.cell(rr_, 1, "D. 校準與份額").font = BOLD; rr_ += 1
rec("k_c（CPU 校準係數）", None, lambda s: f"=PL!D{PB[s]['kc']}", "FY2026 CPU 營收 ÷ 未校準", "接近 1＝法人 CPU 顆數 × 每顆 1,000 美元與 FY2026 CPU 拆分一致", fmt="0.000")
rec("k_o（AI ASIC 校準係數）", None, lambda s: f"=PL!D{PB[s]['ko']}", "FY2026 AI ASIC ÷ 未校準", "< 1：創意 turnkey 只含晶圓與封裝轉售", fmt="0.000")
rec("其他營收毛利率（FY2026 殘差）", None, lambda s: f"=PL!D{PB[s]['gmx']}", "FY2026 毛利殘差", "含 NRE（毛利率高）；FY2025 NRE 約占營收 34%", fmt="0.00%")
rec("CY2028 Broadcom＋Marvell＋聯發科（高）＋創意 其他 ASIC 份額", 1.0, lambda s: f"=Units!E{UR['sum_' + s]}", "avgo、mrvl、mediatek 成品", "應 ≤ 100%", fmt="0.0%", pt=True)
RC.cell(rr_, 1, "E. 預留").font = BOLD; rr_ += 1
RC.cell(rr_, 1, "R4 產業部署路徑（markets/deploy_path/，尚未建立）：上線後以創意 AI 量產營收 ÷ 部署 GW 對照 Tokenomics 非 NVIDIA 加速器比例"); RC.cell(rr_, 10, "預留列：本版無數值"); rr_ += 1
RC.column_dimensions["A"].width = 80
for c in "BCDEFGH": RC.column_dimensions[c].width = 11
RC.column_dimensions["I"].width = 36; RC.column_dimensions["J"].width = 60
RECON_LAST = rr_ - 1

# ---------------- Sources ----------------
SO = wb.create_sheet("Sources"); SO["A1"] = "來源（公司財報與法說內容多為媒體轉述；券商、資料商、媒體皆為 Interested-party）"; SO["A1"].font = BOLD
srcs = [("媒體：經濟日報 2026-07-30（2Q26 營收、營業利益、淨利、EPS、毛利率、turnkey 占 83%；1H26）", UDN2),
        ("媒體：工商時報 2026-07-31（2Q26 EPS；外資 FY2026–28 EPS、目標價 6,500；Visible Alpha 共識）", CT2),
        ("媒體：優分析 2026-09-22（法說轉述：3Q26 營收、毛利率、費用方向；全年毛利率 20% 以上；NRE、turnkey 成長；車用占比）", UAG),
        ("媒體：優分析 2026-01-30（4Q25 法說：不再給季度區間指引；北美 68%；加密貨幣降為個位數）", UA1),
        ("媒體：科技新報（2026 年 9 月營收、前 9 月累計）", SEP26),
        ("媒體：經濟日報 2026-10-07（3Q26 營收 187.5 億；CPU 出貨 2026 約 125 萬、2027 約 400 萬顆；CPU 對 TPU 1:2）", UDN10),
        ("媒體：經濟日報 2026-09-08（CPU 2027 出貨由 250 萬上修至 350 萬顆；HBM4 PHY、COUPE）", UDN9),
        ("媒體：經濟日報 2026-04-28（本國投顧：2027 CSP CPU 50.3%、北美車用 34.9%；每顆 Arm CPU 配 2 顆 ASIC）", UDN4),
        ("媒體：經濟日報 2026-01-31（FY2025 營收 341.4 億；turnkey 約 64%、NRE 約 34%；AI 約 25%）", UDN1),
        ("券商轉述：科技新報 2026-05-05（摩根士丹利：Google CPU 占營收約一半、北美車用 AI 晶片毛利率高於 15%）", MS5),
        ("券商轉述：科技新報 2026-05-08（美銀：1Q26 營收與淨利、Maia 200、Axion 3 奈米）", BOFA5),
        ("資料商：stockanalysis（股數、近 12 月 EPS、前瞻 P/E、共識目標價）", SA), ("資料商：stockanalysis 共識頁（FY2026 共識，2026-09-04）", SAF),
        ("資料商：stockanalysis 財報與比率頁（FY2021–25 損益、ROE、年底 P/E）", SAI + "；" + SAR),
        ("資料商：stockanalysis 股價歷史（10-22 收盤 7,585）", SAH), ("資料商：財報狗（10-22 收盤核對、近四季 EPS）", SDOG),
        ("資料商：stockanalysis 同業前瞻 P/E（3661、2454、AVGO、MRVL）", "https://stockanalysis.com/quote/tpe/3661/；https://stockanalysis.com/quote/tpe/2454/；https://stockanalysis.com/stocks/avgo/；https://stockanalysis.com/stocks/mrvl/"),
        ("資料商：companiesmarketcap 年底 P/E", "https://companiesmarketcap.com/global-unichip-corp/pe-ratio/"),
        ("資料商：digrin 月股價（12 月未調整值）", "https://www.digrin.com/stocks/detail/3443.TW/price"),
        ("資料商：HiStock 歷年 EPS（2018–2025）與季 EPS", HIS),
        ("市場層：tsmc、avgo、mrvl、mediatek 成品與 markets/memory UNITS、SCEN_UNITS（本 repo，唯讀）", f"{rel(TSMC_F)}；{rel(AVGO_F)}；{rel(MRVL_F)}；{rel(MTK_F)}；markets/memory/market.py"),
        ("待補（人工）：3Q26 財報與法說（2026-10-29）；客戶名稱與 CPU、ASIC、車用營收金額；每顆 CPU 營收；HBM 採購方式；下一代 CPU 設計夥伴；1Q、2Q26 稅前利益；2016–2017 EPS 與 2016–2020 淨利率（十年平均）；前瞻 P/E 十年歷史（Bloomberg）", "未找到"),
        ("SemiAnalysis", "未使用")]
for i, (a, b) in enumerate(srcs, 2): SO.cell(i, 1, a); SO.cell(i, 2, b)
SO.column_dimensions["A"].width = 100; SO.column_dimensions["B"].width = 130

# ---------------- README ----------------
RD = wb.create_sheet("README", 0)
lines = [("創意電子（3443.TW）收支與估值模型 v0.1.1（AI 半導體 Project：運算層，ASIC 設計服務＋量產；台積電體系）　2026-10-25", True),
 ("版本：v0.1（初版）。CPU 顆數＝台積電模型的 Google TPU 顆數 × CPU 對 TPU 配比；其他 ASIC 顆數讀台積電模型、每顆核心售價讀 Broadcom 模型，份額與 Broadcom、Marvell、聯發科合計不超過 100%；車用與其他營收用簡化成長。本版未做先驗。", False),
 ("模型命題：創意的營收由 NRE 轉為 Google 自研 Arm CPU（Axion）的 turnkey 量產（晶圓與封裝轉售、毛利率低）。FY2026 營收約翻倍、3Q26 單季 187.5 億元；以 TPU 顆數 × CPU 配比推 CPU 顆數，再加其他 ASIC、車用與 NRE，FY2027–28 的營收與 EPS 有多少？現價約 90 倍前瞻 P/E 需要哪一條路徑？", False),
 ("核心問題：在市場層 TPU 顆數（CY2027 5.5 百萬顆）與 CPU 對 TPU 配比口徑下，創意 FY2027–28 的 CPU 顆數與 EPS 離法人估計（CPU 400 萬顆、FY2027 EPS 94–139 元）多遠；差距靠配比、份額還是車用；現價隱含的倍數與同業差多少。", False), ("", False),
 ("會計年度＝日曆年，FY(t)＝CY(t)，不平移。FY2025A 實績、FY2026 錨點、FY2027–FY2029 驅動因子推導、FY2030–FY2032 收斂法。金額：十億新台幣；EPS：新台幣（基本）。口徑：歸屬母公司淨利（TIFRS）。3Q26 財報 2026-10-29 公布，本版只納入 3Q26 營收實績。", False), ("", False),
 ("驅動因子 → 推導量（全部唯讀上游，不修改）", True),
 ("Google CPU 量產：tsmc 成品的 Google TPU 顆數 × 情境係數（CY2028–29；tsmc／markets/memory 同字母）× CPU 對 TPU 配比（CY2026 0.325 校準；CY2027 起 A／B／C 0.6／0.5／0.4）× 創意份額（CY2028 起 100%／90%／70%）× 每顆營收 1.0K（CY2028 起 ×1.2／1.1／1.0）× k_c（FY2026 以錨點 CPU 拆分 48% 校準）× 匯率 32。", False),
 ("AI ASIC 量產：tsmc 成品的其他 ASIC 顆數 × 創意份額（CY2026 3%；CY2027 6%／4%／2%；CY2028 起 8%／5%／2%）× Broadcom 模型核心售價 6K × k_o。車用：FY2026 錨點（營收 5%）× FY2027 ×4.0／3.0／2.0、FY2028 ×2.0／1.6／1.2。其他（NRE＋其他量產）：FY2027 +15%／+8%／−5%、之後 +10%／+5%／0%。", False),
 ("毛利：AI 量產（CPU、AI ASIC）毛利率 13%（Assumed；情境 A +1pt、C −2pt）、車用 16%；其他營收毛利率為 FY2026 殘差（Derived，含 NRE）後沿用。營業費用彈性 0.3、年增下限 8%；營業外 0.5、稅率 15%、股數 134.01M。", False),
 ("校準：FY2026 ＝ 1H26 實績＋3Q26 營收實績（月營收）＋公司方向性說法（3Q 毛利率較 2Q 下滑、費用持平）＋4Q26 假設（營收季增 35%、毛利率 18.5%）。營收拆分（CPU 48%、AI ASIC 3%、車用 5%）為 Assumed。不用券商數字（只在 Recon 對帳）。", False),
 ("估值：Val 第 0 節目標價＝0.5 × DCF 目標價（股東現金流＝EPS ×（1 − g／ROE），FY2027–32 明確期＋永續，r 10%、g 4%，三情境權重 0.3／0.5／0.2，情境價值負值取 0）＋0.5 × Comparable 目標價（世芯、聯發科、博通、邁威爾前瞻 P/E 中位數 × 12 個月後前瞻 EPS；不含 A 股）；另列 Comparable 單獨目標價、市場共識目標價與差距拆解。十年 P/E 規則值 × FY2027 EPS 保留為參考（不是目標價）。本版未做先驗，無 Andy 合理 P/E 對照。", False),
 ("預估期間：FY2027–FY2029 由驅動因子推導；FY2030 起為收斂法，非驅動因子推導（營收年增由 FY2029 線性收斂到 fade_g、淨利率收斂到 fade_m，FY2032 達到；PL 頁灰底）。十年平均淨利率未找到，B、C 依工作單以 FY2029 × 0.8 代替。", False),
 ("標記：Verified、Interested-party、Analogy、Assumed、Derived。SemiAnalysis 未使用。公司數字多取自媒體轉述財報與法說；券商與資料商為 Interested-party。", False),
 ("", False),
 ("已知限制：(1) 公司未揭露客戶名稱與 CPU、AI ASIC、車用營收金額，FY2026 拆分與毛利率為 Assumed；(2) CPU 顆數綁 TPU 顆數 × 配比，法人 CY2027 約 400 萬顆需配比約 0.73，超出本模型 A 情境 0.6，修訂市場層 TPU 顆數應由市場層提出，不在本模型自改；(3) 其他 ASIC 份額受 Broadcom、Marvell、聯發科份額限制（CY2028 上限 25%），分母含 Ascend 等非台積電產品；(4) 單一客戶集中（Google CPU 約一半營收），情境 C 只反映配比、份額下修，未含失去客戶；(5) 3Q26 毛利率與費用只有方向性說法；1Q、2Q26 稅前利益為推估；(6) 十年 P/E 多為 NRE 主導期；2016–2017 自算 EPS 未找到；(7) R4 產業部署路徑未建（Recon 預留列）。", False), ("", False),
 ("情境對照表（A／B／C 讀同字母的行業情境；公司特有變數 A 最有利、C 最不利）", True),
 ("行業：TPU、其他 ASIC 顆數係數＝tsmc In 頁 um28／um29（＝markets/memory SCEN_UNITS：A 1.0／1.0、B 0.8／0.85、C 0.6／0.6）。不讀 markets/equipment、markets/network（創意無設備、網通營收）。", False),
 ("公司特有：CPU 對 TPU 配比 0.6／0.5／0.4；下一代 CPU 份額 100%／90%／70%、單價係數 1.2／1.1／1.0；其他 ASIC 份額 CY2027 6%／4%／2%、CY2028 起 8%／5%／2%；車用倍數 FY2027 4.0／3.0／2.0、FY2028 2.0／1.6／1.2；其他營收 FY2027 +15%／+8%／−5%、之後 +10%／+5%／0%；量產毛利率調整 +1／0／−2pt。", False)]
for i, (t, b) in enumerate(lines, 1): RD.cell(i, 1, t).font = Font(bold=b)
RD.column_dimensions["A"].width = 170


# ---------------- v0.1.1：上檔空間級距、投資建議、重點觀察（Andy 2026-10-25 14:00、14:17）；各頁只在現有最後一列之後追加，不位移既有列 ----------------
_up = next(r for r in range(1, V.max_row + 1) if V.cell(r, 1).value == "上檔空間：加權目標價")
GRADE_ROW = V.max_row + 2; REC_ROW = GRADE_ROW + 1
V.cell(GRADE_ROW, 1, "上檔空間級距")
V.cell(GRADE_ROW, 2, f'=IF(B{_up}>0.5,"大於 50%",IF(B{_up}>0.25,"25%–50%",IF(B{_up}>=0,"0%–25%",IF(B{_up}>=-0.25,"0%～−25%",IF(B{_up}>=-0.5,"−25%～−50%","小於 −50%")))))')
V.cell(GRADE_ROW, 5, f"以「上檔空間：加權目標價」（B{_up}）分六級（Andy 2026-10-25 14:00）")
V.cell(REC_ROW, 1, "投資建議")
V.cell(REC_ROW, 2, f'=IF(B{_up}>0.5,"強力買進",IF(B{_up}>0.25,"買進",IF(B{_up}>=-0.25,"中立",IF(B{_up}>=-0.5,"放空","強力放空"))))')
V.cell(REC_ROW, 5, "依上檔空間固定規則（Andy 2026-10-25 14:17）：大於 +50% 強力買進、+25%～+50% 買進、−25%～+25% 中立、−50%～−25% 放空、小於 −50% 強力放空")
OBS = ['目標價 NT$2,879（DCF 2,238、Comparable 3,521 各半），比現價 7,585（2026-10-22）低 62.0%，級距 小於 −50%；Comparable 單獨為 −53.6%。主因是倍數不是 EPS：本模型 12 個月後前瞻 EPS 122.95 比共識前瞻 EPS 81.28 高 51.3%，但同業（世芯、聯發科、博通、邁威爾）前瞻 P/E 中位數 28.6 倍，創意本身前瞻 P/E 約 103 倍（2026-10-08）、共識目標價 6,722 隱含 82.7 倍（比同業高 188.8%）；市場給創意的 Google CPU 成長溢價，Comparable 未計入。共識目標價本身也比現價低 11.4%。', '情境分歧最大：FY2028 EPS A／B／C 為 197.6／113.3／50.5，DCF 價格帶 662～3,748；即使 A 情境 DCF 3,748 也只有現價的 49%。', 'CPU 對 TPU 配比是 FY2027 EPS 最敏感的變數：配比 0.6 → 0.4 時 −16.4%、0.75 時 +12.3%；量產毛利率 ±3pt 對 EPS ±9.6%（Sens 頁）。下一個檢驗點是 3Q26 財報與法說（2026-10-29）。']   # 重點觀察（與 README.md 相同文字；數字取自重算後的 Val／PL／Sens 頁）
_rr = RD.max_row + 1
RD.cell(_rr, 1, 'v0.1.1（2026-10-25）：補上檔空間級距、投資建議、重點觀察（數值不變）。Val 頁末新增「上檔空間級距」「投資建議」兩列（Andy 2026-10-25 14:00、14:17 固定規則），README 頁末新增投資建議與重點觀察三點；既有儲存格位置與數值不變。')
RD.cell(_rr + 1, 1, f'="投資建議："&Val!B{REC_ROW}&"（依上檔空間固定規則）"').font = Font(bold=True)
for _j, _t in enumerate(OBS, 1): RD.cell(_rr + 1 + _j, 1, f"重點觀察 {_j}：" + _t)
if os.environ.get("RATING_ROWMAP"):   # 除錯用：輸出新增列號（供撰寫 html.json）；CI 不設此變數
    print("ROWMAP", GRADE_ROW, REC_ROW, _rr, _rr + 1, _rr + 2)


order = ["README", "In", "Q", "Units", "PL", "Val", "Recon", "Sens", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(sys.argv[1])
if os.environ.get("GUC_ROWMAP"):   # 除錯用：輸出列號對照（供撰寫 html.json）；CI 不設此變數
    import json
    json.dump({"PB": PB, "SUMR": {f"{k[0]}|{k[1]}": v for k, v in SUMR.items()}, "UR": UR, "CHK": CHK, "VAL": VAL, "INROW": INROW,
               "SENS_LAST": SENS_LAST, "SENS_NOTE": SENS_NOTE, "RECON_LAST": RECON_LAST, "RULE_ROWS": rule_rows, "RB0": RB0, "WMT_NOTE": WMT_NOTE, "UN_NOTE": UN_NOTE, "T": T, "T10": T10,
               "TPS0": TPS0, "TPS_LAST": TPS_LAST, "PEER_LAST": PEER_LAST, "EPSR": EPSR, "O": O, "CY0": CY0, "R30": R30, "W0": W0},
              open(os.environ["GUC_ROWMAP"], "w"), ensure_ascii=False, indent=1)
